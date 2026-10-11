#!/usr/bin/env python3
"""
arreglos_d9.py — Parrillero-Asador para Eventos 2.2 (SPEC D5, D6, D8, D9, D13): arreglos SÓLO de texto en los 7 ficheros NO
rehechos del producto (6 docx + plantilla de proveedores). Cada arreglo localiza su texto POR CONTENIDO y exige encontrarlo
exactamente las veces esperadas (si no, aborta sin escribir). Parte SIEMPRE de los ficheros publicados de la 1.1
(`git show e7dfc1d8:…`), así que es repetible e idempotente (la huella final no cambia entre ejecuciones). Corre en el VPS
(python-docx + openpyxl). Fork de `paellero-2_2/arreglos_d9.py`.

    /root/venv-guias/bin/python arreglos_d9.py --salida /root/tmp-par-f2/d9     # ensayo: copias corregidas
    /root/venv-guias/bin/python arreglos_d9.py --aplicar                        # escribe en dl/ (paso 3, lo ejecuta el orquestador)

Las cifras de tickets y la tabla de escalado salen de `calculadora_par_es.py` (fuente única con la calculadora).

Sesión Claude Code, 12-oct-2026.
"""
import argparse
import copy
import datetime
import hashlib
import os
import re
import subprocess
import sys
import tempfile
import zipfile

import docx
import docx.text.paragraph
import openpyxl
from docx.shared import Mm
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from openpyxl.worksheet.datavalidation import DataValidation

AQUI = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, AQUI)
REPO = os.path.abspath(os.path.join(AQUI, '..', '..', '..'))
DL = os.path.join(REPO, 'astro-site', 'public', 'dl', 'plan-negocio-parrillero-asador-eventos')
#: Ficheros publicados de la 1.1 (último commit de main antes de la 2.2): el script es idempotente.
REV_ORIGEN = 'e7dfc1d8'
VERSION = '2.2'
FECHA = datetime.datetime(2026, 10, 12)
COMENTARIO = 'aichef.pro/plan-negocio-parrillero-asador-eventos'

import calculadora_par_es as CAL                                                    # noqa: E402

EXP = '10-experiencias-tematicas-eventos.docx'
CARTA = 'carta-12-cortes-carne.docx'
CATALOGO = 'catalogo-equipamiento-utensilios-menaje.docx'
GUIA = 'guia-tipos-parrilla-roadmap-food-truck.docx'
MANUAL = 'manual-tecnico-parrillero-profesional.docx'
CONTRATO = 'modelos-contrato-parrillero-asador-eventos.docx'
PROVEEDORES = 'plantilla-proveedores-parrillero-asador-eventos.xlsx'
DOCX_ORDEN = [EXP, CARTA, CATALOGO, GUIA, MANUAL, CONTRATO]

TITULOS = {
    EXP: '10 Experiencias Temáticas para Eventos Parrillero / Asador',
    CARTA: 'Carta de 12 Cortes de Carne',
    CATALOGO: 'Catálogo de Equipamiento, Utensilios y Menaje Parrillero',
    GUIA: 'Guía de Tipos de Parrilla + Roadmap Food Truck',
    MANUAL: 'Manual Técnico del Parrillero Profesional',
    CONTRATO: 'Modelos de Contrato Parrillero / Asador para Eventos',
}

#: Sustituciones por contenido: fichero → [(viejo, nuevo, nº de párrafos esperados)]. Se aplican en orden, a todos los
#: párrafos (cuerpo y celdas de tabla) y sustituyen la 1.ª aparición dentro de cada párrafo.
TEXTO = {
    EXP: [
        ('Validadas contra mercado real España 2026 — todas tienen demanda confirmada.',
         'La demanda que se indica en cada experiencia es orientativa: compruébala en tu zona antes de apostar por '
         'ella. Los tickets van sin IVA y también son orientativos.', 1),
        ('ajada gallega + salsa villaroy', 'ajada gallega + salsa verde', 1),
        ('Cochinillo segoviano IGP (alternativa)', 'Cochinillo de Segovia (Marca de Garantía; alternativa)', 1),
        ('Lechazo de Castilla DOP en horno de leña', 'Lechazo de Castilla y León IGP en horno de leña', 1),
        ('Cochinillo segoviano IGP', 'Cochinillo de Segovia (Marca de Garantía)', 1),
        ('Verduras a la parrilla (calçots, pimientos, calabacín, alcachofas)',
         'Verduras a la parrilla (pimientos, calabacín y berenjena; alcachofas y calçots solo en su temporada, de '
         'noviembre a abril)', 1),
        ('ensalada payesa + salsa romesco + alioli', 'ensalada payesa + salsa romesco + alioli (hecho con ovoproducto '
                                                      'pasteurizado)', 1),
        ('Provoletas en parrilla pequeña con flambeado en directo',
         'Provoletas en parrilla pequeña con flambeado en directo (extintor a mano, lejos de los invitados y nunca con '
         'alerta de incendios)', 1),
        ('Chorizos al sarmiento con quemado de orujo',
         'Chorizos al sarmiento con quemado de orujo (con las mismas precauciones)', 1),
        ('gastronomic experience', 'experiencia gastronómica', 1),
        ('Las experiencias se combinan. Aquí tres mix probados que funcionan en eventos reales:',
         'Las experiencias se combinan. Aquí tienes tres combinaciones de ejemplo:', 1),
    ],
    CARTA: [
        ('DOP Castilla y León', 'IGP Lechazo de Castilla y León', 1),
        ('IGP Segovia', 'Marca de Garantía Cochinillo de Segovia', 1),
        ('Salsa Villaroy o salsa verde', 'Salsa verde o chimichurri', 1),
        ('Piel a 220 °C últimos 15 min', 'Horno a 220 °C los últimos 15 min', 1),
    ],
    CATALOGO: [
        ('Tres escenarios de inversión (mínimo viable / estándar / premium). Validado contra GGM Gastro España, Brogas '
         '1958, Tienda Biomasa, Texfire y Royal Catering en abril 2026.',
         'Tres escenarios de equipo de fuego y servicio (mínimo viable / estándar / premium). Precios orientativos '
         'consultados en abril de 2026 en GGM Gastro España, Brogas 1958, Tienda Biomasa, Texfire y Royal Catering.', 1),
        ('Parrilla horizontal de 1,2 m con tres parrillas de hierro fundido inox AISI 430.',
         'Parrilla horizontal de 1,2 m con tres parrillas de acero inoxidable AISI 430.', 1),
        ('ROI claro a partir de 4-6 eventos BBQ a 70-90 €/pax.',
         'Se amortiza con los eventos de BBQ que vendas: calcúlalo con el plan financiero del kit.', 1),
        ('Tablas con código de colores alimentario obligatorio en APPCC',
         'Tablas con código de colores alimentario, recomendado en el APPCC (no es obligatorio por ley, pero evita '
         'contaminaciones cruzadas)', 1),
        ('guantes ignífugos certificados EN15025 / ISO6942', 'guantes certificados EN 407 (riesgos térmicos)', 1),
        ('Inversión total', 'Equipo de fuego y servicio (la inversión total está en el plan financiero)', 1),
        ('Mínimo viable + ahumador COMPER-M + asador estaca premium + 4 mesas + furgo adaptada + menaje básico propio '
         '+ indumentaria 2 personas.',
         'Mínimo viable + ahumador COMPER-M + asador estaca premium + 4 mesas + menaje básico propio + indumentaria '
         '2 personas.', 1),
        ('35.000-60.000 €', '30.000-75.000 €', 1),
    ],
    GUIA: [
        ('Tu margen bruto promedio supera el 55 % consistentemente — hay caja para reinversión.',
         'Tu margen bruto real alcanza o supera de forma consistente el margen del plan financiero del kit — hay caja '
         'para reinversión.', 1),
        ('Si tu margen bruto está por debajo del 50 % — antes optimiza pricing.',
         'Si tu margen bruto real está por debajo del margen del plan financiero del kit — antes optimiza el pricing.', 1),
        ('Si no tienes carnet manipulador alimentos avanzado y plan APPCC robusto.',
         'Si no tienes la formación en higiene alimentaria al día y un plan APPCC robusto.', 1),
        ('Eventos finca con licencia OK.', 'La licencia de la finca no levanta la restricción en terreno forestal.', 1),
        ('Bodas finca rural con autorización organizador.',
         'La licencia de la finca no levanta la restricción en terreno forestal: consulta al Consell.', 1),
        ('Autorización ambulante + plan PATRICOVA',
         "Autorización ambulante + Pla especial davant del risc d'incendis forestals", 1),
        ('Autorización municipal + Plan INFOCYL', 'Autorización municipal + Plan INFOCAL', 1),
        ('Llicentzia salmenta ambulanta', 'Autorización municipal de venta ambulante', 1),
        ('MARGEN NETO', 'RESULTADO ANTES DE AMORTIZACIÓN, SUELDO DEL TITULAR E IMPUESTOS (estimación del kit)', 1),
        ('Margen neto %', 'Resultado antes de amortización, sueldo e impuestos (%)', 1),
        ('709-2.000 € (GGM KGE1200, La Vaca Tuerta, Hasbeg)',
         '1.500-3.500 € (parrilla móvil inox de 120-150 cm; precio orientativo, compruébalo con el proveedor)', 1),
        ('Sin licencia conducir camión', 'Sin permiso de conducción de camión', 1),
        ('Carnet B normal', 'Permiso B normal', 1),
        ('Carnet C, peajes', 'Permiso C, peajes', 1),
        ('Concepto principal del 80% de parrilleros eventos en España. Inversión obligatoria.',
         'Es el concepto principal de la mayoría de los parrilleros de eventos. Inversión de base.', 1),
    ],
    MANUAL: [
        ('Temperatura ambiente 30 min antes de cocinar — nunca pieza fría directa al fuego.',
         'Atempera la pieza antes de cocinar: sácala de frío solo el tiempo mínimo (como máximo 30 minutos y a la '
         'sombra, según tu APPCC) — nunca helada y directa al fuego.', 1),
        ('Iniciar parrilla con pieza fría direct al fuego', 'Iniciar la parrilla con la pieza fría directa al fuego', 1),
        ('5. Salsa de mostaza al miel', '5. Salsa de mostaza a la miel', 1),
    ],
    CONTRATO: [
        ('en nombre propio y en su condición de profesional autónomo prestador',
         'en nombre propio o en representación de [razón social del PARRILLERO], prestador', 1),
        (', en adelante el «PARRILLERO» o el «PROVEEDOR».',
         ', actuando en nombre propio o en representación de [razón social del PARRILLERO], en adelante el '
         '«PARRILLERO» o el «PROVEEDOR».', 1),
        ('La factura se emitirá a nombre del CLIENTE en un plazo no superior a 15 días tras el evento.',
         'El PARRILLERO emitirá a nombre del CLIENTE una factura por cada pago recibido, también por la señal (el IVA '
         'se devenga con el cobro, art. 75.Dos de la Ley 37/1992).', 1),
        ('Si la cancelación se produce con más de 60 días de antelación al evento',
         'Si la cancelación se produce con 60 días o más de antelación al evento', 1),
        ('Entre 8 y 29 días de antelación', 'Entre 7 y 29 días de antelación', 1),
        ('Cancelación con más de 30 días de antelación: 20 %', 'Cancelación con 30 días o más de antelación: 20 %', 1),
        ('Cancelación entre 8-14 días: 75 %', 'Cancelación entre 7-14 días: 75 %', 1),
        ('(plan INFOCA, decretos de sequía, alertas forestales)',
         '(planes autonómicos de incendios forestales como el INFOCA, decretos de sequía, alertas forestales, '
         'época de peligro alto)', 1),
        ('En estos supuestos las partes acordarán de buena fe un cambio de fecha o, en su defecto, la devolución '
         'íntegra de las cantidades abonadas.',
         'En estos supuestos, si no se puede cocinar con fuego, las partes intentarán primero el plan B pactado en el '
         'briefing (plancha o parrilla de piedra volcánica a gas, con una botella de menos de 15 kg por aparato, o '
         'equipos sin llama del recinto); si tampoco es posible, acordarán de buena fe un cambio de fecha sin coste '
         'adicional (reprogramación) o, en su defecto, la devolución íntegra de las cantidades abonadas.', 1),
        ('Derechos de acceso, rectificación, supresión, oposición y portabilidad ejercitables',
         'Derechos de acceso, rectificación, supresión, limitación, portabilidad y oposición ejercitables', 1),
        ('se someten a los Juzgados y Tribunales de [Ciudad correspondiente al domicilio del PARRILLERO].',
         'se someten a los Juzgados y Tribunales del domicilio del CLIENTE, por tratarse de un consumidor (art. 90.2 '
         'del Texto Refundido de la Ley General para la Defensa de los Consumidores y Usuarios).', 1),
        ('Retención IRPF aplicable: [7 % primer año, 15 % en adelante] según normativa vigente para autónomos.',
         'Facturación: factura de cada pago, también del anticipo; la factura no lleva retenciones de IRPF (la actividad es '
         'empresarial).',
         1),
        ('seguro de responsabilidad civil profesional vigente con cobertura mínima de [300.000 / 600.000] €',
         'seguro de responsabilidad civil de explotación vigente, con cobertura de incendio y de producto y un mínimo de '
         '[600.000] €', 1),
        ('póliza de RC profesional con cobertura mínima de [600.000] € por siniestro',
         'póliza de RC de explotación, con cobertura de incendio y de producto, y un mínimo de [600.000] € por siniestro',
         1),
        ('Anexo I · Tarifario base B2B 2026', 'Anexo I · Tarifario B2B orientativo: pon tus precios (sin IVA)', 1),
    ],
}

#: tickets de las 10 experiencias (línea bajo «TICKET DE SERVICIO»): viejo → clave de EXP_TICKETS
TICKETS_VIEJOS = [
    ('B2C 39-55 €/pax · B2B planner 75-130 €/pax', '01'), ('B2C 50-70 €/pax · B2B 95-150 €/pax', '02'),
    ('B2C 45-60 €/pax · B2B 80-120 €/pax', '03'), ('B2B 65-95 €/pax · Show cooking 800-1.200 €/jornada', '04'),
    ('B2C 50-65 €/pax · B2B 75-100 €/pax', '05'), ('B2C 55-75 €/pax · B2B 90-140 €/pax', '06'),
    ('B2C 50-70 €/pax · B2B 90-130 €/pax', '07'), ('B2C 35-50 €/pax · B2B 60-85 €/pax', '08'),
    ('Upsell: 8-15 €/pax adicional · Stand-alone: fee 400-800 € evento', '09'),
    ('B2B empresa 35-65 €/pax + fee producción 300-600 €', '10')]
#: «B2C 50-70 €/pax · B2B 95-150 €/pax» aparece una vez por experiencia, salvo 07 que repite el rango de 07 (50-70 / 90-130):
#: cada línea vieja es única en el documento (se comprueba).
COMBOS_VIEJOS = [('Ticket: 95-130 €/pax.', 'boda_premium'), ('Ticket: 65-90 €/pax + fee producción 600-900 €.', 'team_building'),
                 ('Ticket: 80-110 €/pax.', 'boda_costa')]

#: Celdas de tabla: fichero → [(tabla, fila, columna, texto viejo EXACTO, texto nuevo)]
CELDAS = {
    CONTRATO: [
        (0, 0, 1, '€/pax bajo', '€/inv. bajo, sin IVA'), (0, 0, 2, '€/pax medio', '€/inv. medio, sin IVA'),
        (0, 0, 3, '€/pax premium', '€/inv. premium, sin IVA'), (0, 0, 5, 'Fee desplazamiento (€)',
                                                                  'Desplazamiento (€/km, ida y vuelta)'),
        (0, 1, 1, '60', '45'), (0, 1, 2, '90', '60'), (0, 1, 3, '130', '85'), (0, 1, 5, '250', '0,95'),
        (0, 2, 1, '65', '45'), (0, 2, 2, '95', '65'), (0, 2, 3, '150', '90'), (0, 2, 5, '200', '0,95'),
        (0, 3, 1, '45', '26'), (0, 3, 2, '65', '31'), (0, 3, 3, '90', '42'), (0, 3, 5, '350', '0,95'),
        (0, 4, 1, '50', '30'), (0, 4, 2, '75', '40'), (0, 4, 3, '110', '55'), (0, 4, 5, '250', '0,95'),
        (0, 5, 1, 'fee/jornada 600', 'por jornada: a presupuestar'), (0, 5, 2, '850', '—'), (0, 5, 3, '1.200', '—'),
        (0, 5, 5, '200', '0,95'),
        (0, 6, 1, '90', '55'), (0, 6, 2, '130', '75'), (0, 6, 3, '180', '100'), (0, 6, 5, '400', '0,95'),
    ],
}

#: Párrafos NUEVOS: fichero → [(ancla: el párrafo empieza por…, 'despues'|'antes', plantilla: el párrafo empieza por…, texto)]
INSERTAR = {
    CONTRATO: [
        ('Todos los pagos se realizarán mediante transferencia', 'despues',
         'Todos los pagos se realizarán mediante transferencia',
         'Al tratarse de un servicio de comida con fecha determinada, el CLIENTE no dispone del derecho de '
         'desistimiento (art. 103.l del Texto Refundido de la Ley General para la Defensa de los Consumidores y '
         'Usuarios); por eso la política de pagos y cancelaciones de este contrato queda aceptada con la firma.'),
        ('SEXTA · ESTÁNDARES OPERATIVOS Y SLA', 'antes', 'QUINTA · CANCELACIONES B2B', 'QUINTA BIS · FUERZA MAYOR'),
        ('SEXTA · ESTÁNDARES OPERATIVOS Y SLA', 'antes', 'Por su carácter profesional',
         'En caso de fuerza mayor (catástrofe natural, prohibición administrativa, condiciones meteorológicas extremas '
         'declaradas por AEMET, intervención de la autoridad sanitaria o prohibición del uso del fuego por la época o el '
         'nivel de peligro de incendios), las partes intentarán primero el plan B pactado en la Hoja de Pedido (plancha o '
         'parrilla de piedra volcánica a gas, o equipos sin llama del recinto); si no es posible, reprogramarán el '
         'evento sin coste adicional, y si la reprogramación tampoco lo es, el PARRILLERO devolverá íntegramente el '
         'anticipo a la AGENCIA.'),
    ],
}


def _es1(x):
    return ('%.1f' % x).replace('.', ',')


# ---------------------------------------------------------------------------------------------------------------------
# utilidades de docx
# ---------------------------------------------------------------------------------------------------------------------
def _sustituir_parrafo(p, viejo, nuevo):
    """Sustituye la 1.ª aparición de `viejo` conservando el formato: dentro de un run, o fusionando sólo los runs
    afectados (el texto fusionado toma el formato del primero)."""
    for r in p.runs:
        if viejo in r.text:
            r.text = r.text.replace(viejo, nuevo, 1)
            return True
    runs = p.runs
    texto = ''.join(r.text for r in runs)
    i = texto.find(viejo)
    if i < 0:
        return False
    j = i + len(viejo)
    pos = 0
    primero = None
    for r in runs:
        ini, fin = pos, pos + len(r.text)
        pos = fin
        if fin <= i or ini >= j:
            continue
        if primero is None:
            primero = r
            r.text = r.text[:i - ini] + nuevo + (r.text[j - ini:] if fin >= j else '')
        else:
            r.text = r.text[j - ini:] if fin > j else ''
    return primero is not None


def _todos_los_parrafos(d):
    for p in d.paragraphs:
        yield p
    for t in d.tables:
        for row in t.rows:
            for c in row.cells:
                for p in c.paragraphs:
                    yield p


def _clonar(base, texto):
    nuevo = copy.deepcopy(base._p)
    p = docx.text.paragraph.Paragraph(nuevo, base._parent)
    for r in p.runs[1:]:
        r._r.getparent().remove(r._r)
    if p.runs:
        p.runs[0].text = texto
    else:
        p.add_run(texto)
    return nuevo


def _textos(d):
    for p in d.paragraphs:
        yield p.text
    for t in d.tables:
        for row in t.rows:
            for c in row.cells:
                yield c.text
    for s in d.sections:
        for p in list(s.header.paragraphs) + list(s.footer.paragraphs):
            yield p.text


def _buscar(d, prefijo, que):
    cand = [p for p in d.paragraphs if p.text.startswith(prefijo)]
    if len(cand) != 1:
        raise SystemExit('ABORTADO: %s «%s»: %d párrafos (esperado 1)' % (que, prefijo, len(cand)))
    return cand[0]


def _tabla_nueva(d, cabecera, filas, anchos):
    """Tabla con el aspecto de las del manual (estilo, cabecera 1F1B16 con texto blanco, celdas 10 pt)."""
    base = d.tables[3]._tbl
    tbl = copy.deepcopy(base)
    trs = tbl.tr_lst
    tc_h, tc_b = trs[0].tc_lst[0], trs[1].tc_lst[0]
    tr_h, tr_b = copy.deepcopy(trs[0]), copy.deepcopy(trs[1])
    for tr in list(trs):
        tbl.remove(tr)
    grid = tbl.find(qn('w:tblGrid'))
    for g in list(grid):
        grid.remove(g)
    for a in anchos:
        g = OxmlElement('w:gridCol')
        g.set(qn('w:w'), str(a))
        grid.append(g)

    def fila(tr_molde, tc_molde, textos):
        tr = copy.deepcopy(tr_molde)
        for tc in list(tr.tc_lst):
            tr.remove(tc)
        for t, a in zip(textos, anchos):
            tc = copy.deepcopy(tc_molde)
            tc.find(qn('w:tcPr')).find(qn('w:tcW')).set(qn('w:w'), str(a))
            ts = tc.findall('.//' + qn('w:t'))
            ts[0].text = str(t)
            for extra in ts[1:]:
                extra.getparent().remove(extra)
            tr.append(tc)
        return tr
    tbl.append(fila(tr_h, tc_h, cabecera))
    for f in filas:
        tbl.append(fila(tr_b, tc_b, f))
    return tbl


# ---------------------------------------------------------------------------------------------------------------------
# sección nueva del manual
# ---------------------------------------------------------------------------------------------------------------------
def seccion_manual(d, plantillas, ancla):
    """Inserta tras `ancla` (último párrafo del manual) «8. Escalado por invitados»: la tabla de escalado de la
    calculadora y la de recepción y cocción de la carne (SPEC D9)."""
    h, cuerpo, bullet, sub, vacio = plantillas
    cur = ancla._p
    pendientes = []

    def p(base, texto):
        pendientes.append(_clonar(base, texto))

    p(h, '8. Escalado por invitados')
    p(cuerpo, 'La carne se compra por kilos, no por invitados. Esta es la tabla de escalado que usa la calculadora del '
              'kit, para un plato principal de %d g de carne cruda por invitado y servicio en buffet.'
      % CAL.FORMATOS[CAL.REF_FORMATO])
    filas = CAL.tabla_escalado()
    cab = ['Concepto'] + ['%d inv.' % f[0] for f in filas]
    nombres = CAL.ESC_FILAS
    cuerpo_t = []
    for i, rot in enumerate(nombres):
        fila = [rot]
        for f in filas:
            v = f[1 + i]
            fila.append(_es1(v) if i in (0, 1, 5) else '%d' % v)
        cuerpo_t.append(fila)
    pendientes.append(_tabla_nueva(d, cab, cuerpo_t, [3000, 1518, 1518, 1518, 1518]))
    pendientes.append(_clonar(vacio, ''))
    reg = dict((k, v[2]) for k, v in CAL.REGLAS.items())
    ipp, ipa = CAL.SERVICIOS[CAL.REF_SERVICIO]
    for t in (
            'Parrilla de 120 × 60 cm: asa de 12 a 15 kg de carne cruda por tanda; con 2 o 3 tandas da servicio a '
            '60-80 invitados con 400 g cada uno (la tabla toma %s kg por parrilla).' % _es1(reg['kg_parr']),
            'Parrilleros: uno por cada 60-80 invitados en buffet y uno por cada 40-50 con cortes en mesa (la tabla '
            'toma uno por cada %d). Ayudante desde %d invitados (la tabla toma uno por cada %d).'
            % (ipp, reg['ayud_desde'], ipa),
            'Carbón: de 0,7 a 1 kg por kg de carne cruda (la tabla toma %s kg), unos 0,3-0,45 kg por invitado. En el '
            'asado a la cruz, de 2 a 3 kg de leña por kg de carne.' % _es1(reg['carb_kg']),
            'Carne cruda por invitado: de 350 a 450 g con entrantes, de 500 a 600 g si el asado es lo único que se sirve y '
            'de 300 a 400 g en un evento largo. Sin hueso, un 20 % menos; los niños, la mitad.',
            'Merma de cocción: del 25 al 35 %% en vacuno y cerdo a la brasa (la tabla toma el %d %%) y del 40 al 50 %% en '
            'el brisket ahumado.' % round(reg['merma'] * 100),
            'Piezas enteras: un cordero de 10-12 kg en canal da de 15 a 20 raciones y un cochinillo de 4-5 kg, de 6 a 8.',
            'Si cambias una regla en la calculadora, cambia también en esta tabla: son las mismas.'):
        p(bullet, t)
    p(sub, 'Recepción y cocción de la carne')
    p(cuerpo, 'Son las temperaturas del Reglamento (CE) 853/2004, que no obliga al minorista pero es el estándar con el '
              'que recibe la carne y el que adopta tu APPCC, y las de cocción que recomienda AESAN.')
    pendientes.append(_tabla_nueva(
        d, ['Producto', 'Recepción y transporte (máx.)', 'Cocción en el centro'],
        [['Carne fresca (vacuno, cerdo, cordero)', '≤ 7 °C', '70 °C'], ['Despojos (achuras, mollejas)', '≤ 3 °C', '70 °C'],
         ['Aves', '≤ 4 °C', '74 °C'], ['Carne picada', '≤ 2 °C', '70 °C'],
         ['Preparados de carne (chorizo criollo, salchichas, hamburguesas)', '≤ 4 °C', '70 °C']],
        [4200, 2436, 2436]))
    pendientes.append(_clonar(vacio, ''))
    for t in (
            'Transporta y recibe la carne en furgoneta frigorífica con certificado ATP, mide con sonda en la recepción y en '
            'la nevera del evento y guarda el registro.',
            'Mantén en caliente a 63 °C o más, recalienta a 74 °C durante 15 segundos y guarda una comida testigo '
            '(RD 1086/2020, art. 30).',
            'El punto «poco hecho» solo en piezas enteras e intactas de vacuno o cordero y a petición del cliente; nunca en '
            'carne picada, hamburguesa, pollo, embutido fresco ni piezas inyectadas.',
            'Chuletón o T-bone con vértebra: solo de vacuno de un país de riesgo insignificante de EEB (España lo es) o de '
            'menos de 30 meses. Si la columna es material de riesgo, la pieza llega sin ella y el parrillero nunca la '
            'quita: solo un matadero, una sala de despiece o un punto de venta autorizado (RD 1086/2020, art. 26.2). Lee el '
            '«nacido, criado y sacrificado en» de la etiqueta.',
            'Denominación por edad: ternera blanca hasta 8 meses, ternera de 8 a 12, añojo de 12 a 24, novillo o novilla de '
            '24 a 48, y buey, vaca o toro a partir de 48. En la carta va la denominación de la etiqueta: un «chuletón de '
            'ternera» de un añojo es una denominación errónea.',
            '«Madurada X días» solo con el documento del proveedor: la ley no define «madurada» ni sus días. Cómprala ya '
            'madurada; madurar en la cocina de apoyo exige una cámara dedicada.',
            'Guarda las etiquetas y los albaranes del vacuno (Reglamento (CE) 178/2002, art. 18).'):
        p(bullet, t)
    for el in pendientes:
        cur.addnext(el)
        cur = el


# ---------------------------------------------------------------------------------------------------------------------
# docx
# ---------------------------------------------------------------------------------------------------------------------
def docx_(origen, destino, log):
    d = docx.Document(origen)
    f = os.path.basename(destino)
    corto = f[:26]
    # (0) referencias de los párrafos nuevos (se resuelven ANTES de tocar nada)
    ins = []
    for ancla, modo, molde, texto in INSERTAR.get(f, []):
        ins.append((_buscar(d, ancla, 'ancla'), modo, _buscar(d, molde, 'plantilla'), texto))
    if f == MANUAL:
        ps = d.paragraphs
        plant = (_buscar(d, '7. Errores frecuentes', 'plantilla'), _buscar(d, 'Las mermas son el porcentaje', 'plantilla'),
                 _buscar(d, 'No usar sonda digital', 'plantilla'), _buscar(d, 'Chimichurri verde clásico', 'plantilla'),
                 ps[3])
        ancla_manual = ps[-1]
        if not ancla_manual.text.startswith('Ignorar la dirección del viento'):
            raise SystemExit('ABORTADO: el último párrafo del manual ya no es «Ignorar la dirección del viento»')
        if ps[3].text.strip():
            raise SystemExit('ABORTADO: p3 del manual no es un párrafo vacío')
    # (1) sustituciones por contenido
    reglas = list(TEXTO.get(f, []))
    if f == EXP:
        for viejo, cod in TICKETS_VIEJOS:
            reglas.append((viejo, CAL.ticket_texto(cod), 1))
        for viejo, clave in COMBOS_VIEJOS:
            lo, hi = CAL.COMBOS[clave][2]
            reglas.append((viejo, 'Ticket orientativo, sin IVA: %d-%d €/inv.%s' % (
                lo, hi, ' La producción se presupuesta aparte.' if clave == 'team_building' else ''), 1))
    for viejo, nuevo, esperado in reglas:
        n = 0
        for p in _todos_los_parrafos(d):
            c = p.text.count(viejo)
            if c:
                if c != 1 or not _sustituir_parrafo(p, viejo, nuevo):
                    raise SystemExit('ABORTADO: %s: «%s» aparece %d veces en un párrafo o no se pudo sustituir'
                                     % (f, viejo[:60], c))
                n += 1
        if n != esperado:
            raise SystemExit('ABORTADO: %s: «%s» en %d párrafos (esperado %d)' % (f, viejo[:70], n, esperado))
        log.append((f, '%s: «%s» → «%s»' % (corto, viejo[:46], nuevo[:56])))
    # (1b) celdas con texto exacto
    for ti, ri, ci, viejo, nuevo in CELDAS.get(f, []):
        celda = d.tables[ti].rows[ri].cells[ci]
        p = celda.paragraphs[0]
        if celda.text != viejo or not _sustituir_parrafo(p, viejo, nuevo):
            raise SystemExit('ABORTADO: %s t%d.r%d.c%d no es «%s» (es «%s»)' % (f, ti, ri, ci, viejo, celda.text))
        log.append((f, '%s t%d.r%d.c%d: «%s» → «%s»' % (corto, ti, ri, ci, viejo, nuevo)))
    # (2) párrafos nuevos
    for base, modo, molde, texto in ins:
        nuevo = _clonar(molde, texto)
        (base._p.addnext if modo == 'despues' else base._p.addprevious)(nuevo)
        log.append((f, '%s %s «%s»: párrafo nuevo «%s»' % (corto, modo, base.text[:30], texto[:50])))
    if f == MANUAL:
        seccion_manual(d, plant, ancla_manual)
        log.append((f, '%s: sección «8. Escalado por invitados» (tabla de escalado de la calculadora + recepción y '
                       'cocción de la carne)' % corto))
    for s in d.sections:
        s.page_width, s.page_height = Mm(210), Mm(297)
    cp = d.core_properties
    cp.title = '%s · v%s' % (TITULOS[f], VERSION)
    cp.author = cp.last_modified_by = 'AI Chef Pro'
    cp.comments = COMENTARIO
    cp.created = cp.modified = FECHA
    d.save(destino)
    log.append((f, '%s: A4 + docProps (autor, fechas, comments, title v%s)' % (corto, VERSION)))


# ---------------------------------------------------------------------------------------------------------------------
# proveedores
# ---------------------------------------------------------------------------------------------------------------------
NOTA_EUDR = ('Nota EUDR: si tu carbón o tu leña son de marabú (Cuba, África) o de quebracho (Paraguay, Argentina) '
             'importados, pueden ser madera afectada por el Reglamento (UE) 2023/1115 (EUDR). La obligación recae en el '
             'importador desde el 30-12-2026 (empresas grandes y medianas) o el 30-06-2027 (micro y pequeñas): pide a tu '
             'proveedor la declaración de diligencia debida.')


def proveedores(origen, destino, log):
    wb = openpyxl.load_workbook(origen)
    ws0 = wb.worksheets[0]
    assert ws0.title == 'Indice', ws0.title
    assert ws0['A1'].value == 'Plantilla 96 Proveedores · Parrillero / Asador para Eventos', ws0['A1'].value
    assert ws0['C15'].value == 93, ws0['C15'].value
    ws0['A1'] = 'Plantilla de 93 proveedores · Parrillero / Asador para Eventos'
    assert str(ws0['A2'].value).startswith('Listado vivo. Marca como [validado]'), ws0['A2'].value
    ws0['A2'] = ('Listado vivo. «Por verificar» significa que aún no has trabajado con el proveedor ni tienes referencias '
                 'directas: pasa a «Contactado» u «Homologado» cuando las tengas. El score de 1 a 10 lo pones tú según tu '
                 'experiencia.')
    assert str(ws0['A17'].value).startswith('Versión 1.1 · agosto 2026 · '), ws0['A17'].value
    ws0['A17'] = ws0['A17'].value.replace('Versión 1.1 · agosto 2026', 'Versión %s · octubre 2026' % VERSION)
    nval = ninv = nscore = 0
    lista = '"Por verificar,Contactado,Homologado,Investigar"'
    for ws in wb.worksheets[1:]:
        hdr = [r for r in range(1, 8) if ws.cell(r, 1).value == 'Proveedor']
        assert len(hdr) == 1 and ws.cell(hdr[0], 6).value == 'Estado' and ws.cell(hdr[0], 7).value == 'Score 1-10', ws.title
        ultimo = hdr[0]
        for r in range(hdr[0] + 1, ws.max_row + 1):
            if ws.cell(r, 1).value is None:
                continue
            ultimo = r
            est = ws.cell(r, 6).value
            if est == 'validado':
                ws.cell(r, 6).value = 'Por verificar'
                nval += 1
            elif est == 'investigar':
                ws.cell(r, 6).value = 'Investigar'
                ninv += 1
            else:
                raise SystemExit('ABORTADO: estado inesperado %r en %s fila %d' % (est, ws.title, r))
            if ws.cell(r, 7).value:
                ws.cell(r, 7).value = None
                nscore += 1
        dv = DataValidation(type='list', formula1=lista, allow_blank=True)
        dv.add('F%d:F%d' % (hdr[0] + 1, ultimo))
        ws.add_data_validation(dv)
        if ws.title.startswith('2. Carbon'):
            fila = ultimo + 2
            ws.cell(fila, 1).value = NOTA_EUDR
            ws.cell(fila, 1).alignment = openpyxl.styles.Alignment(wrap_text=True, vertical='top')
            ws.merge_cells(start_row=fila, start_column=1, end_row=fila, end_column=8)
            ws.row_dimensions[fila].height = 62
            ws0['D6'] = 'Ver la nota EUDR al final de la hoja'
    assert (nval, ninv, nscore) == (46, 47, 93), (nval, ninv, nscore)
    # pestaña con espacio final
    ren = 0
    for ws in wb.worksheets:
        if ws.title != ws.title.rstrip():
            ws.title = ws.title.rstrip()
            ren += 1
    pr = wb.properties
    assert '96' in pr.title, pr.title
    pr.title = 'Plantilla de 93 proveedores · Parrillero / Asador para Eventos · v%s' % VERSION
    pr.subject = (pr.subject or '').replace('v1.1', 'v' + VERSION)
    pr.creator = pr.lastModifiedBy = 'AI Chef Pro'
    pr.description = COMENTARIO
    pr.created = pr.modified = FECHA
    wb.save(destino)
    log.append((PROVEEDORES, '%s: Índice A1 «96» → 93, A2 sin «validado», A17 versión %s, docProps v%s' % (
        PROVEEDORES[:26], VERSION, VERSION)))
    log.append((PROVEEDORES, '%s: Estado «validado» → «Por verificar» (%d) y «investigar» → «Investigar» (%d), DV de '
                             'Estado en las 10 hojas' % (PROVEEDORES[:26], nval, ninv)))
    log.append((PROVEEDORES, '%s: score prefijado vaciado en %d filas' % (PROVEEDORES[:26], nscore)))
    log.append((PROVEEDORES, '%s: nota EUDR en el bloque del carbón (+ aviso en Índice D6)' % PROVEEDORES[:26]))
    log.append((PROVEEDORES, '%s: pestaña con espacio final renombrada (%d)' % (PROVEEDORES[:26], ren)))


PROHIBIDAS = ['validad', 'freelance', 'carnet', 'carné', 'tarifa plana', 'sin licencia', 'PATRICOVA', 'INFOCYL',
              'retención', 'España 2026', '96 proveedores', 'Proveedores Reales', 'licencia OK', 'villaroy', 'DOP Castilla',
              'IGP Segovia']


def barrido(dest):
    hay = []
    for f in DOCX_ORDEN:
        d = docx.Document(os.path.join(dest, f))
        for t in _textos(d):
            for pal in PROHIBIDAS:
                if pal.lower() in t.lower():
                    hay.append('%s: «%s» en «%s»' % (f[:30], pal, t[:90]))
        for a in ('title', 'subject', 'comments', 'author', 'last_modified_by'):
            v = getattr(d.core_properties, a) or ''
            if 'python-docx' in v:
                hay.append('%s: docProps %s con python-docx' % (f[:30], a))
        if d.core_properties.created.year != 2026:
            hay.append('%s: docProps created %s' % (f[:30], d.core_properties.created))
        for s in d.sections:
            if abs(s.page_width.mm - 210) > 0.5 or abs(s.page_height.mm - 297) > 0.5:
                hay.append('%s: sección no A4 (%.1f × %.1f mm)' % (f[:30], s.page_width.mm, s.page_height.mm))
    wb = openpyxl.load_workbook(os.path.join(dest, PROVEEDORES))
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str):
                    for pal in PROHIBIDAS:
                        if pal.lower() in c.value.lower():
                            hay.append('%s %s: «%s» en «%s»' % (PROVEEDORES[:20], c.coordinate, pal, c.value[:80]))
    pr = wb.properties
    for v in (pr.title, pr.subject):
        if v and ('Reales' in v or '96' in v):
            hay.append('xlsx docProps: %s' % v)
    return hay


def huella(dest):
    """SHA-256 del contenido de todas las partes de los 7 ficheros (ignora fechas del zip): prueba de idempotencia."""
    h = hashlib.sha256()
    for f in DOCX_ORDEN + [PROVEEDORES]:
        with zipfile.ZipFile(os.path.join(dest, f)) as z:
            for n in sorted(z.namelist()):
                dato = z.read(n)
                if n == 'docProps/core.xml':          # openpyxl estampa la fecha de guardado: no es contenido
                    dato = re.sub(rb'(<dcterms:modified[^>]*>)[^<]*', rb'\1', dato)
                h.update(n.encode() + dato)
    return h.hexdigest()[:16]


def main():
    ap = argparse.ArgumentParser()
    g = ap.add_mutually_exclusive_group(required=True)
    g.add_argument('--salida')
    g.add_argument('--aplicar', action='store_true')
    a = ap.parse_args()
    dest = DL if a.aplicar else a.salida
    os.makedirs(dest, exist_ok=True)
    log = []
    tmp = tempfile.mkdtemp(prefix='d9-origen-')
    for f in DOCX_ORDEN + [PROVEEDORES]:
        rel = os.path.relpath(os.path.join(DL, f), REPO)
        with open(os.path.join(tmp, f), 'wb') as fh:
            fh.write(subprocess.check_output(['git', '-C', REPO, 'show', '%s:%s' % (REV_ORIGEN, rel)]))
    for f in DOCX_ORDEN:
        docx_(os.path.join(tmp, f), os.path.join(dest, f), log)
    proveedores(os.path.join(tmp, PROVEEDORES), os.path.join(dest, PROVEEDORES), log)
    print('\n'.join(l for _, l in log))
    por = {}
    for f, _ in log:
        por[f] = por.get(f, 0) + 1
    for f, n in por.items():
        print('  %3d  %s' % (n, f))
    print('D9: %d arreglos en %d ficheros → %s' % (len(log), len(DOCX_ORDEN) + 1, dest))
    print('HUELLA %s' % huella(dest))
    hay = barrido(dest)
    if hay:
        print('BARRIDO: %d coincidencias prohibidas' % len(hay))
        print('\n'.join(hay))
        return 1
    print('BARRIDO: 0 coincidencias prohibidas')
    return 0


if __name__ == '__main__':
    sys.exit(main())
