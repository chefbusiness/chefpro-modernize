#!/usr/bin/env python3
"""
semilla_par.py — Parrillero / Asador ES 2.2 (SPEC D1/D3/D10): prepara el ORIGEN que `planes-v2_0/main.py` convierte en el plan
financiero y el checklist del parrillero. Copia adaptada de `paellero-2_2/semilla_pae.py`.

Copia los dos xlsx del Food Truck ES v2.2 publicados (`dl/plan-negocio-food-truck/`, SOLO LECTURA) con los nombres de
fichero del parrillero y escribe lo que el módulo de contenido no alcanza (F1-camino §1):
  · plan: títulos de hoja heredados del fichero (Inversión A1/A2, PyG A2, Punto Equilibrio A2), las 14 partidas de
    inversión (filas 7-20: rótulo, importe, nota, ¿IVA?) y los costes fijos propios del P&L (filas 26-35);
  · checklist: título, pestañas, cabeceras de fase y las 83 tareas de `datos_checklist.py` (12/16/14/10/16/15) con el
    pie «Tareas completadas:» debajo; DV, formato condicional y contadores los rehace el motor (`checklist_ok_y_contador`).
Corre en el VPS (openpyxl). Después: `main.py --producto plan-negocio-parrillero-asador-eventos --dry-run --origen <salida>`.

    /root/venv-guias/bin/python semilla_par.py --salida /root/tmp-par-f2/motor/origen

Sesión Claude Code, 12-oct-2026.
"""
import argparse
import copy
import os
import shutil

import importlib
import sys

import openpyxl

import datos_checklist as DC

AQUI = os.path.dirname(os.path.abspath(__file__))
REPO = os.path.abspath(os.path.join(AQUI, '..', '..', '..'))
FT = os.path.join(REPO, 'astro-site', 'public', 'dl', 'plan-negocio-food-truck')
PLAN_FT, CHECK_FT = 'plan-financiero-food-truck.xlsx', 'checklist-apertura-food-truck.xlsx'
PLAN, CHECK = 'plan-financiero-parrillero-asador-eventos.xlsx', 'checklist-apertura-parrillero-asador-eventos.xlsx'
PID = 'plan-negocio-parrillero-asador-eventos'

TITULOS_PLAN = {
    ('Inversión Inicial', 'A1'): 'PLAN DE NEGOCIO: PARRILLERO / ASADOR PARA EVENTOS',
    ('Inversión Inicial', 'A2'): 'Inversión inicial estimada — parrilla y asado en directo para eventos',
    ('PyG 3 Años', 'A2'): 'Escenario realista — parrillero / asador para eventos',
    ('Punto Equilibrio', 'A2'): 'Punto de equilibrio — parrillero / asador para eventos',
}

#: Filas 7-20 de 'Inversión Inicial' (molde FT): (rótulo, importe, nota, ¿lleva IVA?). Research §3: 49.900 € en 15
#: partidas; «web, fotos, marca y dossier» (2.500) y «lanzamiento» (2.000) van juntas para caber en las 14 filas.
#: Sin palabras de recuento en el rótulo (motor.RX_RECUENTO) ni «registro» fuera de la fila de trámites (AMORTIZABLE).
INVERSION = [
    ('Furgoneta frigorífica usada con certificado ATP', 20000,
     'Furgoneta usada con caja frigorífica y certificado ATP vigente (RD 237/2000) para llevar la carne a 7 °C o '
     'menos. Estimación del kit (18.000-21.950 € en el mercado de ocasión)', 'Sí'),
    ('Rotulación, estanterías y anclajes', 1200,
     'Rotulación con la marca, estanterías con amarres y anclajes para llevar en pie y sujetas las botellas de gas y '
     'los sacos de carbón', 'Sí'),
    ('Remolque-parrilla homologado con cruz', 7000,
     'Remolque de hasta 750 kg con parrilla y soporte de cruz: la pieza de las bodas y las fiestas. Estimación del '
     'kit (6.000-9.000 €); no hay precio público de referencia', 'Sí'),
    ('Parrillas móviles inox de 120 cm y brasero', 3200,
     'Dos parrillas móviles de carbón y leña de 120 cm (cada una sirve 60-80 invitados en dos o tres tandas) y un '
     'brasero para tener brasa de reserva', 'Sí'),
    ('Cruces, kit de costillar y ahumador pequeño', 1800,
     'Cuatro cruces para cordero o cochinillo, el kit de costillar y un ahumador offset pequeño para el BBQ texano',
     'Sí'),
    ('Plan B de gas y seguridad contra incendios', 2200,
     'Plancha o parrilla de piedra volcánica de gas con sus botellas y reguladores (una botella de menos de 15 kg '
     'por aparato), tres extintores ABC y uno de CO2, mantas ignífugas, contenedor metálico de cenizas y EPI '
     'térmicos', 'Sí'),
    ('Frío: armario refrigerado portátil, isotermos, acumuladores y sondas', 1800,
     'Cadena de frío de la carne desde la cocina de apoyo hasta el evento, y sondas para las temperaturas del APPCC',
     'Sí'),
    ('Carpa, mesas, lavamanos portátil y protector de césped', 1800,
     'Zona de parrilla bajo carpa (a 3 m o más de la parrilla), mesa de corte y de servicio, lavamanos para el '
     'personal y protector del césped o del pavimento', 'Sí'),
    ('Cuchillería, tablas de colores, básculas, GN y uniformes', 1900,
     'Cuchillos de deshuese y de trinchar, tablas por colores, básculas, cubetas GN y uniformes de la plantilla',
     'Sí'),
    ('Marca, web, fotos, dossier y lanzamiento', 4500,
     'Identidad, web con formulario de presupuesto, sesión de fotos y vídeo de la brasa, dossier para bodas y '
     'empresas, y degustaciones de lanzamiento', 'Sí'),
    ('Software, tablet y datáfono', 800, 'Presupuestos, agenda y facturación conforme a Veri*factu', 'Sí'),
    ('Constitución SL + notaría + registro', 1000,
     'SL unipersonal. Las tasas y el arancel del Registro no llevan IVA; los honorarios del notario y de la '
     'gestoría sí: cambia la columna de al lado a «Sí» si tu presupuesto es sobre todo de honorarios', 'No'),
    ('Declaración responsable, plan APPCC, alérgenos y formación', 1200,
     'Redacción del plan APPCC, fichas de alérgenos de cada menú, formación en higiene alimentaria del equipo y '
     'trámites sanitarios', 'Sí'),
    ('Stock inicial de producto', 1500,
     'Primer pedido de carbón y leña, especias, salsas, desechables compostables y la primera carga de gas', 'Sí'),
]
assert len(INVERSION) == 14 and sum(x[1] for x in INVERSION) == 49900

#: Filas 26-35 de 'PyG 3 Años' (costes fijos propios). Research §3: 11.050 €/año en 10 partidas (10.920 tras la demo), una por fila del
#: molde (no se suprimen filas: desplazaría el P&L, catering F2-NOTAS §3.4). Sin «eventos», «clientes» ni «días»
#: en el rótulo (motor.RX_RECUENTO). Las claves de `contenido….FIJOS` son estos rótulos normalizados.
FIJOS = {
    26: ('Asesoría fiscal y laboral (nóminas de fijos discontinuos)', 1800),
    27: ('Furgoneta y remolque: mantenimiento, ITV e inspección ATP', 1400),
    28: ('Marketing y publicidad', 4500),
    29: ('Software de gestión comercial, CRM y teléfono', 900),
    30: ('Mantenimiento de extintores y revisión del gas', 170),   # research 300: −130 € por la demo §2.11.5 (F2-NOTAS §3)
    31: ('Uniformes, EPI y lavandería', 500),
    32: ('Gestión de residuos (cenizas, grasa y aceite)', 250),
    33: ('Prevención de riesgos laborales y vigilancia de la salud', 400),
    34: ('Degustaciones comerciales y pruebas de menú', 800),
    35: ('Afilado y reposición de cuchillería', 200),
}
assert sum(x[1] for x in FIJOS.values()) == 10920


def _supuestos(wb):
    """Los valores de `0. Supuestos` del caso parrillero. `grupo_a` CONSERVA lo que el libro ya trae en las celdas
    verdes (§1.3: dato del cliente) y sólo usa el valor del contenido en una celda vacía: partiendo del libro del
    FT, sin este paso el caso seguiría siendo el del food truck (45 × 12 € × 250, medido en el primer dry-run)."""
    p2 = os.path.join(os.path.dirname(AQUI), 'planes-v2_0')
    sys.path.insert(0, p2)
    import grupo_a
    contenido = importlib.import_module('contenido_plan_negocio_parrillero_asador_eventos.a')
    coords = dict((t[0], t[1]) for t in grupo_a.SUPUESTOS_BASE)
    ws = wb['0. Supuestos']
    escritas = 0
    for clave, t in contenido.SUPUESTOS.items():
        if t[2] is None:
            continue
        assert clave in coords, clave
        ws[coords[clave]] = t[2]
        escritas += 1
    return escritas


def _docprops(wb):
    """REVISION-FINAL B1: el motor no reescribe `description` ni `keywords` y heredaba los del FT."""
    wb.properties.description = 'aichef.pro/' + PID
    wb.properties.keywords = 'plan negocio parrillero asador eventos, AI Chef Pro'


def plan(salida):
    destino = os.path.join(salida, PLAN)
    shutil.copy2(os.path.join(FT, PLAN_FT), destino)
    wb = openpyxl.load_workbook(destino)
    print('  supuestos escritos:', _supuestos(wb))
    _docprops(wb)
    for (hoja, coord), texto in TITULOS_PLAN.items():
        wb[hoja][coord] = texto
    ws = wb['Inversión Inicial']
    for i, (rot, imp, nota, iva) in enumerate(INVERSION):
        r = 7 + i
        assert not str(ws.cell(r, 1).value).lower().startswith(('fianza', 'alquiler', 'imprevistos')), r
        ws.cell(r, 1).value, ws.cell(r, 2).value, ws.cell(r, 4).value, ws.cell(r, 5).value = rot, imp, nota, iva
    ws = wb['PyG 3 Años']
    assert str(ws['A25'].value).startswith('Seguros') and str(ws['A36'].value).startswith('Amortización')
    for r, (rot, imp) in FIJOS.items():
        ws.cell(r, 1).value, ws.cell(r, 2).value = rot, imp
        ws.cell(r, 6).value = None          # la nota la pone el contenido (FIJOS)
    wb.save(destino)
    return destino


def _copiar_estilo(origen, destino):
    destino.font, destino.border = copy.copy(origen.font), copy.copy(origen.border)
    destino.fill, destino.alignment = copy.copy(origen.fill), copy.copy(origen.alignment)
    destino.number_format, destino.protection = origen.number_format, copy.copy(origen.protection)


def checklist(salida):
    destino = os.path.join(salida, CHECK)
    shutil.copy2(os.path.join(FT, CHECK_FT), destino)
    wb = openpyxl.load_workbook(destino)
    fases = [ws for ws in wb.worksheets if ws.title[:2] in ('F1', 'F2', 'F3', 'F4', 'F5', 'F6')]
    assert len(fases) == 6
    for ws, (titulo, cabecera, tareas) in zip(fases, DC.FASES):
        ws.title = titulo
        ws['A1'], ws['A2'] = DC.TITULO, cabecera
        assert ws['B4'].value and 'TRÁMITE' in ws['B4'].value.upper(), ws['B4'].value
        modelo = [copy.copy(ws.cell(5, c)) for c in range(1, 6)]
        alto = ws.row_dimensions[5].height
        pie = [copy.copy(ws.cell(r, c)) for r in range(5, ws.max_row + 1) for c in range(1, 6)
               if isinstance(ws.cell(r, c).value, str) and 'completad' in ws.cell(r, c).value.lower()]
        assert pie, ws.title
        fila_pie = pie[0].row
        estilo_pie = [copy.copy(ws.cell(fila_pie, c)) for c in range(1, 6)]
        # vaciar la zona de tareas y el pie (valores, DV y formato condicional los rehace el motor)
        for r in range(5, ws.max_row + 1):
            for c in range(1, 6):
                ws.cell(r, c).value = None
            ws.row_dimensions[r].height = None
        ws.data_validations.dataValidation = []
        ws.conditional_formatting = type(ws.conditional_formatting)()
        for i, (tarea, resp, plazo, nota) in enumerate(tareas):
            r = 5 + i
            for c, v in enumerate(('☐', tarea, resp, plazo, nota), start=1):
                cel = ws.cell(r, c)
                _copiar_estilo(modelo[c - 1], cel)
                cel.value = v
            ws.row_dimensions[r].height = alto
        r = 5 + len(tareas) + 1
        for c in range(1, 6):
            _copiar_estilo(estilo_pie[c - 1], ws.cell(r, c))
        ws.cell(r, 2).value = 'Tareas completadas:'
        ws.cell(r, 3).value = '=COUNTIF(A5:A%d,"✓")' % (4 + len(tareas))
        ws.cell(r, 4).value = 'de'
        ws.cell(r, 5).value = '=COUNTIF(B5:B%d,"?*")-COUNTIF(A5:A%d,"N/A")' % (4 + len(tareas), 4 + len(tareas))
    # B1: `motor.asegurar_instrucciones` no reescribe la hoja si ya existe: la del FT nombraba su producto
    ins = wb['Instrucciones']
    assert 'plan-negocio-food-truck' in str(ins['A3'].value), ins['A3'].value
    ins['A3'] = 'Este fichero forma parte del producto «%s» de AI Chef Pro.' % PID
    _docprops(wb)
    wb.save(destino)
    return destino


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--salida', required=True)
    a = ap.parse_args()
    if os.path.isdir(a.salida):
        shutil.rmtree(a.salida)
    os.makedirs(a.salida)
    print('plan      →', plan(a.salida))
    print('checklist →', checklist(a.salida), '(%d tareas: %s)' % (
        sum(len(t) for _, _, t in DC.FASES), '/'.join(str(len(t)) for _, _, t in DC.FASES)))


if __name__ == '__main__':
    main()
