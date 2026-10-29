#!/usr/bin/env python3
"""
calculadora_taq_es.py — Plan de Negocio: Taquería ES 1.0: `calculadora-trompo-escandallo-taqueria.xlsx`.
Hojas `Trompo` · `Tortilla` · `Escandallo` · `Instrucciones`. Gemelo de código: `plan-pizzeria/calculadora_piz_es.py`
(helpers, inject_cache, autotest con pycel y un `calcular()` en Python, protección sin contraseña, verdes desbloqueadas, A4,
docProps AI Chef Pro). Todos los valores por defecto son «estimación del kit, editable» (F1-research-taqueria.md §2 y §8).
  · Trompo, bloque A (el cono al pastor): kilos crudos, adobo, merma de cocción y goteo, recortes, sobrante y g por taco →
    marinados, servidos, rendimiento, tacos por cono y por kg crudo, €/kg servido, carne por taco, kg crudos para N tacos.
  · Trompo, bloque B (otras proteínas): 6 guisos con kg crudos, €/kg, rendimiento tras la cocción y g por taco.
  · Tortilla: comprada frente a nixtamal propio → € por tortilla, tortillas por kg y EQUILIBRIO en tortillas al día.
  · Escandallo: 19 ingredientes × 10 columnas (8 tacos + orden de 3 al pastor + gringa) → coste, food cost con semáforo,
    margen, PVP sugerido y margen por canal (sala · para llevar · plataforma) de la orden y la gringa.
Valores que otros scripts leen de la CACHÉ del xlsx (dict `EXPUESTOS`, función `valores_expuestos(ruta)`; celdas por defecto):
    TACOS_POR_CONO           Trompo!B%(TACOS_POR_CONO)s      ≈ 277,7 tacos
    RENDIMIENTO_TROMPO       Trompo!B%(RENDIMIENTO_TROMPO)s      0,648 (64,8 %)
    COSTE_TACO_PASTOR        Escandallo!D%(COSTE_TACO_PASTOR)s     ≈ 0,61 € (coste con merma)
    FC_TACO_PASTOR           Escandallo!D%(FC_TACO_PASTOR)s     ≈ 22,5 %
    EQUILIBRIO_TORTILLAS_DIA Tortilla!B%(EQUILIBRIO_TORTILLAS_DIA)s     ≈ 697 tortillas al día
Corre en el VPS (openpyxl + pycel):

    /root/venv-guias/bin/python calculadora_taq_es.py --salida <dir>      # genera + inject_cache + autotest
    /root/venv-guias/bin/python calculadora_taq_es.py --salida <dir> --autotest

Sesión Claude Code, 29-oct-2026.
"""
import argparse
import math
import os
import re
import sys
from collections import OrderedDict

import openpyxl
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill, Protection
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

AQUI = os.path.dirname(os.path.abspath(__file__))
PD = os.path.dirname(AQUI)
sys.path.insert(0, AQUI)
import taq_version                                                                 # noqa: E402,F401  (parchea motor: 1.0 · noviembre 2026)
import motor                                                                       # noqa: E402

PID = 'plan-negocio-taqueria'
FICHERO = 'calculadora-trompo-escandallo-taqueria.xlsx'
TITULO = 'Calculadora del trompo y escandallo por taco'
HT, HO, HE, HI = 'Trompo', 'Tortilla', 'Escandallo', 'Instrucciones'
VERDE = PatternFill('solid', fgColor='E8F5E9')
GRIS = PatternFill('solid', fgColor='F2F2F2')
ROJO = PatternFill('solid', bgColor='F8CBAD', fgColor='F8CBAD')
AMBAR = PatternFill('solid', bgColor='FFE699', fgColor='FFE699')
OK = PatternFill('solid', bgColor='C6E0B4', fgColor='C6E0B4')
FUENTE = 'Calibri'
EUR, ENT, PCT, DEC, GR, EUR3 = '#,##0.00 €', '#,##0', '0.0%', '#,##0.0', '#,##0.0', '#,##0.000 €'
MODOS = ['Tortilla comprada', 'Nixtamal propio']
FC_OK, FC_VIG, FC_ALTO = 'En objetivo', 'Vigila el coste', 'Por encima del objetivo'
SEM_COMPRA, SEM_PROPIA = 'Compra la tortilla', 'Hazla tú'
EST = ' Estimación del kit, editable.'

# ---------------------------------------------------------------------------------------------------------------------
# Datos (F1-research §2 y §8). Importes sin IVA salvo el PVP de carta.
# ---------------------------------------------------------------------------------------------------------------------
#: hoja Trompo, bloque A: (clave, rótulo, formato, valor por defecto, constante gris, nota)
ENT_A = [
    ('K', 'Carne cruda montada en el cono (kg)', DEC, 15.0, False, 'Pesa la carne antes del adobo, ya cortada y lista para montar.' + EST),
    ('a', 'Adobo sobre la carne (%)', PCT, 0.08, False, 'Achiote, chiles, especias y piña del adobo, sobre el peso de la carne.' + EST),
    ('Pc', 'Precio de la carne cruda (€/kg, sin IVA)', EUR, 5.50, False, 'Cabeza de lomo o paleta de cerdo. Pon el de tu proveedor.' + EST),
    ('Pa', 'Precio del adobo (€/kg, sin IVA)', EUR, 6.00, False, 'Pon el de tu proveedor o el coste de tu receta.' + EST),
    ('pk', 'Piña de la punta del cono (kg)', DEC, 0.0, False, 'Por defecto 0: la piña de cada taco se cuenta aparte en la hoja «Escandallo». Si pones kilos, suman al coste del cono.'),
    ('pp', 'Precio de la piña (€/kg, sin IVA)', EUR, 2.00, False, 'Solo cuenta si pones kilos de piña arriba.' + EST),
    ('mc', 'Merma de cocción y goteo (%)', PCT, 0.35, False, 'Lo que pierde el cono al asarse: grasa y agua. No hay una cifra publicada fiable: mídela con tu primer cono (hoja «Instrucciones»).' + EST),
    ('mr', 'Recortes no servibles (%)', PCT, 0.05, False, 'Puntas quemadas, grasa dura y nervios que no se sirven.' + EST),
    ('S', 'Sobrante al cierre (kg)', DEC, 0.0, False, 'Carne cocinada que queda sin servir al cerrar. Pésala; por defecto 0.'),
    ('r', 'Parte del sobrante que se reaprovecha (%)', PCT, 0.0, False, 'Por ejemplo en un guiso del día siguiente. La parte que se tira resta tacos.'),
    ('gt', 'Carne servida por taco (g)', GR, 35.0, False, 'Gramos de carne cocinada y cortada en cada taco.' + EST),
    ('N', 'Número de tacos que quieres preparar', ENT, 100, False, 'Número entero mayor que 0. Te dice cuánta carne cruda necesitas.'),
    ('KG', 'Gramos por kilo (constante)', ENT, 1000, True, 'Constante: un kilo son 1.000 gramos. No se edita.'),
    ('ref25', 'Referencia: tacos por kg crudo en un puesto de CDMX', ENT, 25, True,
     'Referencia fija: un trompo de 50 a 100 kg da entre 1.250 y 2.500 tacos, unos 25 por kg crudo (Milenio). No se edita.'),
]
SAL_A = [
    ('marin', 'Carne marinada (kg)', DEC, '{K}*(1+{a})', 'Carne cruda más adobo.'),
    ('servidos', 'Carne servida del cono (kg)', '#,##0.00', '{marin}*(1-{mc}-{mr})-{S}*(1-{r})',
     'Marinada menos la merma y los recortes, menos el sobrante que se tira.'),
    ('rend', 'RENDIMIENTO DEL TROMPO (%)', PCT, '{servidos}/{K}', 'Servida ÷ carne cruda montada.'),
    ('tacos', 'TACOS QUE SALEN DEL CONO', ENT, '{servidos}*{KG}/{gt}', 'Servida × 1.000 ÷ gramos por taco.'),
    ('tacos_kg', 'Tacos por kg de carne cruda', DEC, '{tacos}/{K}', 'Compárala con la referencia de CDMX de arriba.'),
    ('vs_ref', 'Tus tacos por kg frente a la referencia de CDMX (%)', '0%', '{tacos_kg}/{ref25}', 'Por debajo de 100 %: sirves más gramos por taco o pierdes más.'),
    ('coste_cono', 'Coste del cono (€)', EUR, '{K}*{Pc}+{K}*{a}*{Pa}+{pk}*{pp}', 'Carne + adobo + piña de la punta.'),
    ('coste_kg', 'COSTE POR KG SERVIDO (€)', EUR, '{coste_cono}/{servidos}', 'Coste del cono ÷ kg servidos. Alimenta el escandallo.'),
    ('carne_taco', 'COSTE DE CARNE POR TACO (€)', EUR3, '{coste_kg}*{gt}/{KG}', 'Coste por kg servido × gramos por taco ÷ 1.000.'),
    ('kgN', 'Carne cruda para preparar los N tacos (kg)', '#,##0.00', '{N}*{gt}/({KG}*{rend})', 'N × gramos por taco ÷ (1.000 × rendimiento).'),
]
#: bloque B: (guiso, kg crudos, €/kg crudo, rendimiento tras la cocción, g por taco)
GUISOS = [
    ('Carnitas', 10.0, 5.20, 0.60, 35.0),
    ('Suadero', 10.0, 7.50, 0.55, 35.0),
    ('Cochinita pibil', 10.0, 5.50, 0.60, 40.0),
    ('Barbacoa de res', 10.0, 7.00, 0.50, 40.0),
    ('Tinga de pollo', 10.0, 4.00, 0.65, 35.0),
    ('Campechano', 10.0, 7.00, 0.55, 35.0),
]
NB = len(GUISOS)

#: hoja Tortilla: (clave, rótulo, formato, valor, nota)
ENT_T = [
    ('tp', 'Precio del paquete de tortillas (€, sin IVA)', EUR, 14.00, 'Tortilla de maíz de 12 cm. El minorista cuesta unos 0,21 € la unidad; el profesional, un 30 % menos.' + EST),
    ('tn', 'Piezas por paquete', ENT, 100, 'Número entero mayor que 0.' + EST),
    ('tg', 'Peso de cada tortilla (g)', GR, 22.0, 'Una tortilla de 12 cm pesa unos 22 g.' + EST),
    ('mz', 'Precio del maíz para nixtamal (€/kg, sin IVA)', EUR, 1.20, 'Maíz blanco en grano. Pon el de tu proveedor.' + EST),
    ('mm', 'Masa que sale de 1 kg de maíz (kg)', '#,##0.00', 2.0, 'El maíz nixtamalizado absorbe agua: de 1 kg de grano salen unos 2 kg de masa.' + EST),
    ('cg', 'Cal, gas, luz y agua del nixtamal (€ por kg de maíz)', EUR, 2.00, 'Todo lo que se gasta en cocer, moler y prensar, por kilo de maíz.' + EST),
    ('gn', 'Masa por tortilla propia (g)', GR, 22.0, 'Masa cruda que lleva cada tortilla antes de prensarla.' + EST),
    ('tw', 'Coste anual del tortillero (€, con cargas)', EUR, 18900.0, 'Unas tres cuartas partes de una jornada, con seguridad social.' + EST),
    ('ti', 'Inversión en molino, prensa y mesa (€, sin IVA)', EUR, 18000.0, 'Molino de nixtamal, tortilladora y mesa de trabajo.' + EST),
    ('ty', 'Años de vida de la inversión', ENT, 6, 'La inversión se reparte en estos años.' + EST),
    ('td', 'Días de apertura al año', ENT, 300, 'Los del caso del plan de negocio.' + EST),
    ('tpd', 'Tortillas que gastas al día', ENT, 280, 'Las del caso del plan de negocio. Cámbialas por las tuyas.' + EST),
]
SAL_T = [
    ('c_comp', 'Coste por tortilla comprada (€)', EUR3, '{tp}/{tn}', 'Paquete ÷ piezas.'),
    ('tort_kg', 'Tortillas que salen de 1 kg comprado', DEC, '{KG}/{tg}', 'Para pedir al proveedor: 1.000 ÷ gramos por pieza.'),
    ('maiz_t', 'Maíz por tortilla propia (g)', GR, '{gn}/{mm}', 'Masa por tortilla ÷ masa que sale de 1 kg de maíz.'),
    ('tort_maiz', 'Tortillas que salen de 1 kg de maíz', DEC, '{mm}*{KG}/{gn}', ''),
    ('c_prop', 'Coste por tortilla propia (€, sin la persona ni la inversión)', EUR3, '{maiz_t}*({mz}+{cg})/{KG}', 'Maíz + cal, gas, luz y agua.'),
    ('ahorro', 'Ahorro por tortilla al hacerla tú (€)', EUR3, '{c_comp}-{c_prop}', 'Si sale cero o negativo, comprar siempre gana.'),
    ('fijo_dia', 'Coste fijo diario de hacerlas tú (€)', EUR, '({tw}+{ti}/{ty})/{td}', 'Persona + inversión repartida, por día de apertura.'),
    ('eq', 'EQUILIBRIO (tortillas al día)', ENT, 'IF({ahorro}>0,({tw}+{ti}/{ty})/({ahorro}*{td}),0)',
     'Con más tortillas al día que esta cifra, hacerlas compensa. Con menos, compra.'),
    ('dia_comp', 'Coste diario con tortilla comprada (€)', EUR, '{tpd}*{c_comp}', 'Para las tortillas al día que has puesto.'),
    ('dia_prop', 'Coste diario con tortilla propia (€)', EUR, '{tpd}*{c_prop}+{fijo_dia}', 'Incluye la persona y la inversión.'),
    ('sem', 'Qué te conviene', 'General', 'IF(AND({ahorro}>0,{tpd}>={eq}),"%s","%s")' % (SEM_PROPIA, SEM_COMPRA), 'Según tus tortillas al día frente al equilibrio.'),
]

#: hoja Escandallo
ESC_IN = [
    ('iva', 'IVA de la comida (%)', PCT, 0.10, 'Comida preparada, en sala, para llevar y a domicilio: 10 % (Ley 37/1992, art. 91.Uno.1.1.º y 2.2.º).'),
    ('obj', 'Food cost objetivo (% sobre el precio sin IVA)', PCT, 0.25, 'Referencia de taquería: 22-30 % en el taco al pastor.' + EST),
    ('tol', 'Margen de vigilancia sobre el objetivo (puntos)', PCT, 0.05, 'Hasta este margen por encima del objetivo, el semáforo dice «Vigila el coste».'),
    ('merm', 'Merma de producción sobre el coste (%)', PCT, 0.05, 'Salsas crudas que se descartan, tortillas rotas, cebolla y cilantro que no se aprovechan.' + EST),
    ('modo', 'Tortilla de maíz que usa el escandallo', 'General', MODOS[0], 'Elige el coste por tortilla de la hoja «Tortilla»: comprada o nixtamal propio.'),
    ('env', 'Envase por pedido para llevar (€, sin IVA)', EUR, 0.50, 'Bolsa, caja y tarrinas de salsa de un pedido.' + EST),
    ('com', 'Comisión de la plataforma de reparto (%)', PCT, 0.30, 'Sobre el precio sin IVA del pedido. Uber Eats ronda el 30 %; revisa tu contrato.' + EST),
]
#: gramos: (ingrediente, precio €/kg o enlace 'cono'/'bN', origen del precio)
ING_G = [
    ('Carne al pastor (del cono)', 'cono', 'enlazado desde la hoja Trompo (€/kg servido)'),
    ('Carnitas', 'b0', 'enlazado desde la hoja Trompo (€/kg cocinado)'),
    ('Suadero', 'b1', 'enlazado desde la hoja Trompo (€/kg cocinado)'),
    ('Cochinita pibil', 'b2', 'enlazado desde la hoja Trompo (€/kg cocinado)'),
    ('Barbacoa de res', 'b3', 'enlazado desde la hoja Trompo (€/kg cocinado)'),
    ('Tinga de pollo', 'b4', 'enlazado desde la hoja Trompo (€/kg cocinado)'),
    ('Campechano', 'b5', 'enlazado desde la hoja Trompo (€/kg cocinado)'),
    ('Cebolla picada', 1.20, 'estimación del kit'),
    ('Cilantro', 12.00, 'estimación del kit'),
    ('Piña', 2.00, 'estimación del kit'),
    ('Lima', 2.50, 'estimación del kit'),
    ('Salsa roja', 3.00, 'estimación del kit'),
    ('Salsa verde', 3.00, 'estimación del kit'),
    ('Salsa macha', 9.00, 'estimación del kit'),
    ('Queso Oaxaca o chihuahua', 8.50, 'estimación del kit'),
    ('Guacamole', 11.00, 'estimación del kit (depende del aguacate)'),
    ('Frijoles de la olla', 3.50, 'estimación del kit'),
]
#: unidades: (ingrediente, precio €/ud o None si es el de la hoja Tortilla, origen)
ING_U = [
    ('Tortilla de maíz (€/ud)', None, 'enlazado desde la hoja Tortilla, según el modo elegido'),
    ('Tortilla de harina (€/ud)', 0.20, 'estimación del kit'),
]
TM, TH = 'Tortilla de maíz (€/ud)', 'Tortilla de harina (€/ud)'
PAST, CAR, SUA, COC, BAR, TIN, CAM = [g[0] for g in ING_G[:7]]
#: columna → (PVP con IVA, {ingrediente en g: g}, {tortillas: uds})
RECETAS = OrderedDict([
    ('Taco al pastor', (3.00, {PAST: 35, 'Cebolla picada': 10, 'Cilantro': 3, 'Piña': 8, 'Lima': 5, 'Salsa roja': 8, 'Salsa verde': 7}, {TM: 1})),
    ('Taco de carnitas', (3.00, {CAR: 35, 'Cebolla picada': 10, 'Cilantro': 3, 'Lima': 5, 'Salsa roja': 7, 'Salsa verde': 8}, {TM: 1})),
    ('Taco de suadero', (3.00, {SUA: 35, 'Cebolla picada': 10, 'Cilantro': 3, 'Lima': 5, 'Salsa roja': 7, 'Salsa macha': 4, 'Salsa verde': 4}, {TM: 1})),
    ('Taco de cochinita', (3.50, {COC: 40, 'Cebolla picada': 10, 'Lima': 5, 'Salsa verde': 10}, {TM: 1})),
    ('Taco de barbacoa', (3.50, {BAR: 40, 'Cebolla picada': 10, 'Cilantro': 3, 'Lima': 5, 'Salsa roja': 10}, {TM: 1})),
    ('Taco de tinga', (3.00, {TIN: 35, 'Cebolla picada': 5, 'Cilantro': 2, 'Queso Oaxaca o chihuahua': 10, 'Salsa roja': 8, 'Salsa verde': 5}, {TM: 1})),
    ('Taco campechano', (3.50, {CAM: 35, 'Cebolla picada': 10, 'Cilantro': 3, 'Lima': 5, 'Salsa roja': 8, 'Salsa verde': 7}, {TM: 1})),
    ('Taco vegetal', (3.50, {'Frijoles de la olla': 40, 'Queso Oaxaca o chihuahua': 20, 'Guacamole': 25, 'Cebolla picada': 5, 'Cilantro': 3,
                            'Lima': 5, 'Salsa verde': 10}, {TM: 1})),
    ('Orden de 3 al pastor', (8.50, {PAST: 105, 'Cebolla picada': 30, 'Cilantro': 9, 'Piña': 24, 'Lima': 15, 'Salsa roja': 24, 'Salsa verde': 21}, {TM: 3})),
    ('Gringa al pastor', (7.50, {PAST: 70, 'Queso Oaxaca o chihuahua': 60, 'Piña': 15, 'Cebolla picada': 10, 'Cilantro': 3, 'Salsa roja': 12,
                                'Salsa verde': 8}, {TH: 1})),
])
NP, NG, NU = len(RECETAS), len(ING_G), len(ING_U)
PCOL = [get_column_letter(4 + j) for j in range(NP)]                              # D..M
J_PASTOR, J_ORDEN, J_GRINGA = 0, NP - 2, NP - 1

# ---------------------------------------------------------------------------------------------------------------------
# Disposición de las hojas (filas)
# ---------------------------------------------------------------------------------------------------------------------
M = OrderedDict()                                                                  # hoja Trompo, entradas A → fila
M_HA = 5
for _i, (_k, *_r) in enumerate(ENT_A):
    M[_k] = M_HA + 1 + _i
M_RA = M_HA + 1 + len(ENT_A) + 1
S_A = OrderedDict((k, M_RA + 1 + i) for i, (k, *_r) in enumerate(SAL_A))
M_HB = M_RA + 1 + len(SAL_A) + 1
B0 = M_HB + 1                                                                      # primera fila de guisos
B_NOTA = B0 + NB

T = OrderedDict()                                                                  # hoja Tortilla
_f_ = 5
for _k, *_r in ENT_T:
    _f_ += 1
    if _k in ('mz', 'tpd'):
        _f_ += 1                                                                   # fila de cabecera de sección
    T[_k] = _f_
T_RES = _f_ + 2
S_T = OrderedDict((k, T_RES + 1 + i) for i, (k, *_r) in enumerate(SAL_T))

E = OrderedDict((k, 5 + i) for i, (k, *_r) in enumerate(ESC_IN))
E_CAB = 5 + len(ESC_IN) + 1
E_PVP = E_CAB + 1
E_G0 = E_PVP + 1
E_G1 = E_G0 + NG - 1
E_UH = E_G1 + 1
E_U0 = E_UH + 1
E_U1 = E_U0 + NU - 1
E_R = E_U1 + 2
RES = [
    ('peso', 'Peso servido en el plato, sin tortilla (g)', GR),
    ('c_mat', 'Coste de materia prima (€)', EUR3),
    ('c_tot', 'COSTE CON MERMA (€)', EUR3),
    ('pvp_sin', 'Precio de venta sin IVA (€)', EUR),
    ('fc', 'FOOD COST (%)', PCT),
    ('sem', 'Semáforo contra el objetivo', 'General'),
    ('margen', 'Margen bruto (€, precio sin IVA − coste)', EUR),
    ('pvp_sug', 'PVP con IVA sugerido para llegar al objetivo (€)', EUR),
]
R = OrderedDict((k, E_R + 1 + i) for i, (k, *_r) in enumerate(RES))
E_CH = E_R + len(RES) + 2
CANAL = [
    ('sala_e', 'Sala: margen (€)', EUR), ('sala_p', 'Sala: margen (%)', PCT),
    ('llev_e', 'Para llevar: margen (€, descontado el envase)', EUR), ('llev_p', 'Para llevar: margen (%)', PCT),
    ('plat_e', 'Plataforma: margen (€, descontados envase y comisión)', EUR), ('plat_p', 'Plataforma: margen (%)', PCT),
]
CH = OrderedDict((k, E_CH + 1 + i) for i, (k, *_r) in enumerate(CANAL))
CH_COL = {J_ORDEN: 'D', J_GRINGA: 'E'}                                             # columna del bloque de canal → receta
E_FIN = E_CH + len(CANAL) + 2

EXPUESTOS = OrderedDict([
    ('TACOS_POR_CONO', (HT, 'B%d' % S_A['tacos'])),
    ('RENDIMIENTO_TROMPO', (HT, 'B%d' % S_A['rend'])),
    ('COSTE_TACO_PASTOR', (HE, '%s%d' % (PCOL[J_PASTOR], R['c_tot']))),
    ('FC_TACO_PASTOR', (HE, '%s%d' % (PCOL[J_PASTOR], R['fc']))),
    ('EQUILIBRIO_TORTILLAS_DIA', (HO, 'B%d' % S_T['eq'])),
])
for _n, (_h, _c) in EXPUESTOS.items():
    __doc__ = __doc__.replace('%(' + _n + ')s', re.sub(r'^[A-Z]+', '', _c))


def valores_defecto():
    """Todas las entradas por defecto en un solo dict plano (claves = las de las celdas verdes)."""
    v = OrderedDict()
    for k, _r, _f2, val, _c, _n in ENT_A:
        v[k] = val
    for i, (_n, kg, p, rd, g) in enumerate(GUISOS):
        v['bk%d' % i], v['bp%d' % i], v['br%d' % i], v['bg%d' % i] = kg, p, rd, g
    for k, _r, _f2, val, _n in ENT_T:
        v[k] = val
    for k, _r, _f2, val, _n in ESC_IN:
        v[k] = val
    for i, (_n, p, _o) in enumerate(ING_G):
        if not isinstance(p, str):
            v['p%d' % i] = p
    v['th'] = ING_U[1][1]
    for j, (_nom, (pvp, gr, ud)) in enumerate(RECETAS.items()):
        v['pvp%d' % j] = pvp
        for i, (ni, *_r) in enumerate(ING_G):
            v['g%d_%d' % (i, j)] = float(gr.get(ni, 0))
        for i, (ni, *_r) in enumerate(ING_U):
            v['u%d_%d' % (i, j)] = float(ud.get(ni, 0))
    return v


# ---------------------------------------------------------------------------------------------------------------------
# Cálculo en Python (espejo de las fórmulas) — fuente del autotest
# ---------------------------------------------------------------------------------------------------------------------
def calcular(v):
    r = OrderedDict()
    r['marin'] = v['K'] * (1 + v['a'])
    r['servidos'] = r['marin'] * (1 - v['mc'] - v['mr']) - v['S'] * (1 - v['r'])
    r['rend'] = r['servidos'] / v['K']
    r['tacos'] = r['servidos'] * v['KG'] / v['gt']
    r['tacos_kg'] = r['tacos'] / v['K']
    r['vs_ref'] = r['tacos_kg'] / v['ref25']
    r['coste_cono'] = v['K'] * v['Pc'] + v['K'] * v['a'] * v['Pa'] + v['pk'] * v['pp']
    r['coste_kg'] = r['coste_cono'] / r['servidos']
    r['carne_taco'] = r['coste_kg'] * v['gt'] / v['KG']
    r['kgN'] = v['N'] * v['gt'] / (v['KG'] * r['rend'])
    r['B'] = []
    for i in range(NB):
        b = OrderedDict()
        b['coc'] = v['bk%d' % i] * v['br%d' % i]
        b['ckg'] = v['bp%d' % i] / v['br%d' % i]
        b['tacos'] = b['coc'] * v['KG'] / v['bg%d' % i]
        b['ctaco'] = b['ckg'] * v['bg%d' % i] / v['KG']
        r['B'].append(b)
    t = OrderedDict()
    t['c_comp'] = v['tp'] / v['tn']
    t['tort_kg'] = v['KG'] / v['tg']
    t['maiz_t'] = v['gn'] / v['mm']
    t['tort_maiz'] = v['mm'] * v['KG'] / v['gn']
    t['c_prop'] = t['maiz_t'] * (v['mz'] + v['cg']) / v['KG']
    t['ahorro'] = t['c_comp'] - t['c_prop']
    t['fijo_dia'] = (v['tw'] + v['ti'] / v['ty']) / v['td']
    t['eq'] = (v['tw'] + v['ti'] / v['ty']) / (t['ahorro'] * v['td']) if t['ahorro'] > 0 else 0
    t['dia_comp'] = v['tpd'] * t['c_comp']
    t['dia_prop'] = v['tpd'] * t['c_prop'] + t['fijo_dia']
    t['sem'] = SEM_PROPIA if (t['ahorro'] > 0 and v['tpd'] >= t['eq']) else SEM_COMPRA
    r['T'] = t
    pg = []
    for i, (_n, p, _o) in enumerate(ING_G):
        pg.append(r['coste_kg'] if p == 'cono' else r['B'][int(p[1:])]['ckg'] if isinstance(p, str) else v['p%d' % i])
    r['pg'] = pg
    r['pu'] = [t['c_comp'] if v['modo'] == MODOS[0] else t['c_prop'], v['th']]
    for j in range(NP):
        gs = [v['g%d_%d' % (i, j)] for i in range(NG)]
        us = [v['u%d_%d' % (i, j)] for i in range(NU)]
        p = OrderedDict()
        p['peso'] = sum(gs)
        p['c_mat'] = sum(g_ * q for g_, q in zip(gs, pg)) / v['KG'] + sum(u_ * q for u_, q in zip(us, r['pu']))
        p['c_tot'] = p['c_mat'] * (1 + v['merm'])
        p['pvp_sin'] = v['pvp%d' % j] / (1 + v['iva'])
        p['fc'] = p['c_tot'] / p['pvp_sin']
        p['sem'] = FC_OK if p['fc'] <= v['obj'] else (FC_VIG if p['fc'] <= v['obj'] + v['tol'] else FC_ALTO)
        p['margen'] = p['pvp_sin'] - p['c_tot']
        p['pvp_sug'] = p['c_tot'] / v['obj'] * (1 + v['iva'])
        p['sala_e'] = p['margen']
        p['sala_p'] = p['sala_e'] / p['pvp_sin']
        p['llev_e'] = p['sala_e'] - v['env']
        p['llev_p'] = p['llev_e'] / p['pvp_sin']
        p['plat_e'] = p['llev_e'] - v['com'] * p['pvp_sin']
        p['plat_p'] = p['plat_e'] / p['pvp_sin']
        r['p%d' % j] = p
    return r


# ---------------------------------------------------------------------------------------------------------------------
# Helpers (de la Pizzería)
# ---------------------------------------------------------------------------------------------------------------------
def _f(ws, coord, v, b=False, sz=11, i=False, fmt=None, wrap=False, fill=None, locked=True):
    c = ws[coord]
    c.value = v
    c.font = Font(name=FUENTE, size=sz, bold=b, italic=i)
    if fmt:
        c.number_format = fmt
    c.alignment = Alignment(wrap_text=wrap, vertical='top')
    if fill:
        c.fill = fill
    c.protection = Protection(locked=locked)
    return c


def _nota(ws, fila, texto, desde='E', hasta='F', chars=80.0):
    if not texto:
        return
    _f(ws, '%s%d' % (desde, fila), texto, wrap=True)
    if desde != hasta:
        ws.merge_cells('%s%d:%s%d' % (desde, fila, hasta, fila))
    ws.row_dimensions[fila].height = max(16, 15 * math.ceil(len(texto) / chars))


def _imprimir(ws, horizontal=False):
    ws.page_setup.paperSize = 9                                                   # A4
    ws.page_setup.orientation = 'landscape' if horizontal else 'portrait'
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.oddFooter.center.text = 'aichef.pro/%s · página &P de &N' % PID
    ws.print_options.horizontalCentered = True
    ws.protection.sheet = True                                                    # sin contraseña
    ws.sheet_view.showGridLines = False


def _dv(ws, tipo, celdas, f1=None, f2=None, op=None, titulo='Valor no válido', msg=''):
    dv = DataValidation(type=tipo, formula1=f1, formula2=f2, operator=op, allow_blank=False, showErrorMessage=True,
                        errorTitle=titulo, error=msg)
    ws.add_data_validation(dv)
    for c in celdas:
        dv.add(c)


def a_excel(expr, ref):
    return re.sub(r'\{(\w+)\}', lambda m: ref[m.group(1)], expr)


def _kg_ext():
    return "'%s'!$B$%d" % (HT, M['KG'])


# ---------------------------------------------------------------------------------------------------------------------
# Hojas
# ---------------------------------------------------------------------------------------------------------------------
def hoja_trompo(wb):
    ws = wb.active
    ws.title = HT
    for col, w in (('A', 54), ('B', 15), ('C', 15), ('D', 16), ('E', 14), ('F', 14), ('G', 14), ('H', 12), ('I', 15)):
        ws.column_dimensions[col].width = w
    _f(ws, 'A1', 'CALCULADORA DEL TROMPO — CUÁNTOS TACOS SALEN DEL CONO Y CUÁNTO CUESTA LA CARNE', b=True, sz=14)
    _f(ws, 'A2', 'Escribe solo en las celdas VERDES: todo lo demás se calcula. Importes sin IVA. Todos los valores por defecto son una '
                 'estimación del kit, editable: cámbialos por los de tu proveedor y, sobre todo, por lo que midas en tu propio cono.', wrap=True)
    ws.merge_cells('A2:I2')
    ws.row_dimensions[2].height = 32
    _f(ws, 'A3', 'AVISO FIJO. La merma de cocción y goteo (35 %) no tiene una fuente publicada fiable: es una estimación del kit. Pesa tu '
                 'primer cono crudo, lo servido y el sobrante, y ajusta la merma hasta que el rendimiento coincida con lo que mides.',
       wrap=True, b=True, fill=GRIS)
    ws.merge_cells('A3:I3')
    ws.row_dimensions[3].height = 36
    _f(ws, 'A%d' % M_HA, 'A. EL CONO AL PASTOR', b=True, fill=GRIS)
    _f(ws, 'B%d' % M_HA, 'Valor', b=True, fill=GRIS)
    _f(ws, 'C%d' % M_HA, 'Notas', b=True, fill=GRIS)
    for k, rot, fmt, val, const, nota in ENT_A:
        f = M[k]
        _f(ws, 'A%d' % f, rot)
        if const:
            _f(ws, 'B%d' % f, val, fmt=fmt, fill=GRIS)
        else:
            _f(ws, 'B%d' % f, val, fmt=fmt, fill=VERDE, locked=False)
        _nota(ws, f, nota, desde='C', hasta='I', chars=105.0)
    ref = dict((k, '$B$%d' % f) for k, f in M.items())
    ref.update((k, '$B$%d' % f) for k, f in S_A.items())
    for col, t in zip('ABC', ('RESULTADO DEL CONO', 'Valor', 'Notas')):
        _f(ws, '%s%d' % (col, M_RA), t, b=True, fill=GRIS)
    for k, rot, fmt, expr, nota in SAL_A:
        f = S_A[k]
        neg = k in ('rend', 'tacos', 'coste_kg', 'carne_taco')
        _f(ws, 'A%d' % f, rot, b=neg)
        _f(ws, 'B%d' % f, '=' + a_excel(expr, ref), fmt=fmt, b=neg)
        _nota(ws, f, nota, desde='C', hasta='I', chars=105.0)
    cab = ('B. OTRAS PROTEÍNAS (guisos y asados)', 'kg crudos', '€/kg crudo (sin IVA)', 'Rendimiento tras la cocción (%)', 'g por taco',
           'kg cocinados', '€/kg cocinado', 'Tacos', '€ de carne por taco')
    for i, t in enumerate(cab):
        _f(ws, '%s%d' % (get_column_letter(1 + i), M_HB), t, b=True, fill=GRIS, wrap=True)
    ws.row_dimensions[M_HB].height = 46
    for i, (nom, kg, p, rd, g) in enumerate(GUISOS):
        f = B0 + i
        _f(ws, 'A%d' % f, nom, fill=VERDE, locked=False)
        _f(ws, 'B%d' % f, kg, fmt=DEC, fill=VERDE, locked=False)
        _f(ws, 'C%d' % f, p, fmt=EUR, fill=VERDE, locked=False)
        _f(ws, 'D%d' % f, rd, fmt='0%', fill=VERDE, locked=False)
        _f(ws, 'E%d' % f, g, fmt=GR, fill=VERDE, locked=False)
        _f(ws, 'F%d' % f, '=B%d*D%d' % (f, f), fmt='#,##0.00')
        _f(ws, 'G%d' % f, '=C%d/D%d' % (f, f), fmt=EUR)
        _f(ws, 'H%d' % f, "=F%d*%s/E%d" % (f, '$B$%d' % M['KG'], f), fmt=ENT)
        _f(ws, 'I%d' % f, "=G%d*E%d/%s" % (f, f, '$B$%d' % M['KG']), fmt=EUR3, b=True)
    _f(ws, 'A%d' % B_NOTA, 'Los guisos de cocción larga (carnitas, barbacoa, cochinita) rinden menos que el cono: el rendimiento y el precio de '
                           'cada fila son una estimación del kit, editable. El €/kg cocinado de cada fila alimenta la hoja «Escandallo».',
       i=True, wrap=True)
    ws.merge_cells('A%d:I%d' % (B_NOTA, B_NOTA))
    ws.row_dimensions[B_NOTA].height = 34
    cel = lambda k: 'B%d' % M[k]
    _dv(ws, 'decimal', [cel(k) for k in ('K', 'Pc', 'Pa', 'gt')], f1='0', op='greaterThan', msg='Escribe un número mayor que 0.')
    _dv(ws, 'decimal', [cel(k) for k in ('pk', 'pp', 'S')], f1='0', op='greaterThanOrEqual', msg='Escribe un número mayor o igual que 0.')
    _dv(ws, 'decimal', [cel(k) for k in ('a', 'r')], f1='0', f2='1', op='between', titulo='Porcentaje no válido',
        msg='Escribe un porcentaje entre 0 y 1 (0,08 = 8 %).')
    _dv(ws, 'decimal', [cel('mc')], f1='0', f2='0.6', op='between', titulo='Merma no válida', msg='Escribe entre 0 y 0,6 (0,35 = 35 %).')
    _dv(ws, 'decimal', [cel('mr')], f1='0', f2='0.3', op='between', titulo='Recortes no válidos', msg='Escribe entre 0 y 0,3 (0,05 = 5 %).')
    _dv(ws, 'whole', [cel('N')], f1='1', f2='1000000', op='between', titulo='Número no válido', msg='Escribe un número entero de 1 o más.')
    rows = range(B0, B0 + NB)
    _dv(ws, 'decimal', ['%s%d' % (c, f) for c in 'BCE' for f in rows], f1='0', op='greaterThan', msg='Escribe un número mayor que 0.')
    _dv(ws, 'decimal', ['D%d' % f for f in rows], f1='0.05', f2='1', op='between', titulo='Rendimiento no válido',
        msg='Escribe entre 0,05 y 1 (0,6 = 60 %).')
    _imprimir(ws, horizontal=True)


def hoja_tortilla(wb):
    ws = wb.create_sheet(HO)
    for col, w in (('A', 56), ('B', 17), ('C', 70)):
        ws.column_dimensions[col].width = w
    _f(ws, 'A1', 'TORTILLA — COMPRADA FRENTE A NIXTAMAL PROPIO Y EQUILIBRIO EN TORTILLAS AL DÍA', b=True, sz=14)
    _f(ws, 'A2', 'Escribe solo en las celdas VERDES. Importes sin IVA. Los valores por defecto son una estimación del kit, editable: cámbialos '
                 'por los de tu proveedor. A volúmenes de taquería de barrio, comprar suele ganar; esta hoja te dice dónde cambia la cuenta.',
       wrap=True)
    ws.merge_cells('A2:C2')
    ws.row_dimensions[2].height = 34
    ref = dict((k, '$B$%d' % f) for k, f in T.items())
    ref.update((k, '$B$%d' % f) for k, f in S_T.items())
    ref['KG'] = _kg_ext()
    sec = {'tp': 'A. TORTILLA COMPRADA', 'mz': 'B. NIXTAMAL PROPIO (maíz en grano, molino y prensa)', 'tpd': 'C. TU CASO'}
    for k, rot, fmt, val, nota in ENT_T:
        f = T[k]
        if k in sec:
            _f(ws, 'A%d' % (f - 1), sec[k], b=True, fill=GRIS)
            _f(ws, 'B%d' % (f - 1), 'Valor', b=True, fill=GRIS)
            _f(ws, 'C%d' % (f - 1), 'Notas', b=True, fill=GRIS)
        _f(ws, 'A%d' % f, rot)
        _f(ws, 'B%d' % f, val, fmt=fmt, fill=VERDE, locked=False)
        _nota(ws, f, nota, desde='C', hasta='C', chars=78.0)
    for col, t in zip('ABC', ('RESULTADO', 'Valor', 'Notas')):
        _f(ws, '%s%d' % (col, T_RES), t, b=True, fill=GRIS)
    for k, rot, fmt, expr, nota in SAL_T:
        f = S_T[k]
        neg = k in ('c_comp', 'c_prop', 'eq', 'sem')
        _f(ws, 'A%d' % f, rot, b=neg)
        _f(ws, 'B%d' % f, '=' + a_excel(expr, ref), fmt=fmt, b=neg, wrap=(k == 'sem'))
        _nota(ws, f, nota, desde='C', hasta='C', chars=78.0)
    ws.row_dimensions[S_T['sem']].height = 32
    sem = '$B$%d' % S_T['sem']
    ws.conditional_formatting.add('B%d' % S_T['sem'], FormulaRule(formula=['%s="%s"' % (sem, SEM_PROPIA)], fill=OK))
    ws.conditional_formatting.add('B%d' % S_T['sem'], FormulaRule(formula=['%s="%s"' % (sem, SEM_COMPRA)], fill=AMBAR))
    cel = lambda k: 'B%d' % T[k]
    _dv(ws, 'decimal', [cel(k) for k in ('tp', 'tg', 'mz', 'mm', 'gn')], f1='0', op='greaterThan', msg='Escribe un número mayor que 0.')
    _dv(ws, 'decimal', [cel(k) for k in ('cg', 'tw', 'ti')], f1='0', op='greaterThanOrEqual', msg='Escribe un número mayor o igual que 0.')
    _dv(ws, 'whole', [cel(k) for k in ('tn', 'ty', 'tpd')], f1='1', f2='1000000', op='between', titulo='Número no válido',
        msg='Escribe un número entero de 1 o más.')
    _dv(ws, 'whole', [cel('td')], f1='1', f2='366', op='between', titulo='Días no válidos', msg='Escribe un número entero de 1 a 366.')
    _imprimir(ws)


def hoja_escandallo(wb):
    ws = wb.create_sheet(HE)
    for col, w in (('A', 40), ('B', 14), ('C', 7)):
        ws.column_dimensions[col].width = w
    for col in PCOL:
        ws.column_dimensions[col].width = 14
    ws.column_dimensions['N'].width = 44
    _f(ws, 'A1', 'ESCANDALLO POR TACO — 8 TACOS, UNA ORDEN Y UNA GRINGA DE EJEMPLO, EN GRAMOS', b=True, sz=14)
    _f(ws, 'A2', 'Escribe solo en las celdas VERDES (precios, gramos, nombres, tortillas y precios de carta). Importes sin IVA salvo el precio de '
                 'carta. Los precios son una estimación del kit, editable: pon los de tu proveedor. Las carnes y la tortilla de maíz vienen de '
                 'las hojas «%s» y «%s».' % (HT, HO), wrap=True)
    ws.merge_cells('A2:N2')
    ws.row_dimensions[2].height = 32
    for col, t in zip('ABD', ('DATOS GENERALES', 'Valor', 'Notas')):
        _f(ws, '%s4' % col, t, b=True, fill=GRIS)
    for k, rot, fmt, val, nota in ESC_IN:
        f = E[k]
        _f(ws, 'A%d' % f, rot)
        _f(ws, 'B%d' % f, val, fmt=fmt, fill=VERDE, locked=False, wrap=(k == 'modo'))
        _nota(ws, f, nota, desde='D', hasta='N', chars=150.0)
    ws.row_dimensions[E['modo']].height = 32
    for col, t in (('A', 'INGREDIENTE (en gramos)'), ('B', '€/kg sin IVA'), ('C', 'Unidad')):
        _f(ws, '%s%d' % (col, E_CAB), t, b=True, fill=GRIS, wrap=True)
    _f(ws, 'N%d' % E_CAB, 'ORIGEN DEL PRECIO', b=True, fill=GRIS, wrap=True)
    _f(ws, 'A%d' % E_PVP, 'Precio de carta con IVA (€)', b=True)
    for j, (nom, (pvp, _g, _u)) in enumerate(RECETAS.items()):
        col = PCOL[j]
        _f(ws, '%s%d' % (col, E_CAB), nom, b=True, fill=VERDE, wrap=True, locked=False)
        _f(ws, '%s%d' % (col, E_PVP), pvp, fmt=EUR, fill=VERDE, locked=False, b=True)
    ws.row_dimensions[E_CAB].height = 32
    for i, (nom, p, origen) in enumerate(ING_G):
        f = E_G0 + i
        if isinstance(p, str):
            _f(ws, 'A%d' % f, nom)
            if p == 'cono':
                _f(ws, 'B%d' % f, "='%s'!$B$%d" % (HT, S_A['coste_kg']), fmt=EUR)
            else:
                _f(ws, 'B%d' % f, "='%s'!$G$%d" % (HT, B0 + int(p[1:])), fmt=EUR)
        else:
            _f(ws, 'A%d' % f, nom, fill=VERDE, locked=False)
            _f(ws, 'B%d' % f, p, fmt=EUR, fill=VERDE, locked=False)
        _f(ws, 'C%d' % f, 'g')
        for j, (_n, (_pvp, gr, _ud)) in enumerate(RECETAS.items()):
            _f(ws, '%s%d' % (PCOL[j], f), float(gr.get(nom, 0)), fmt='#,##0.0;-#,##0.0;"-"', fill=VERDE, locked=False)
        _f(ws, 'N%d' % f, origen, i=True)
    _f(ws, 'A%d' % E_UH, 'INGREDIENTE (en unidades)', b=True, fill=GRIS)
    _f(ws, 'B%d' % E_UH, '€/ud sin IVA', b=True, fill=GRIS)
    _f(ws, 'C%d' % E_UH, 'Unidad', b=True, fill=GRIS)
    for i, (nom, p, origen) in enumerate(ING_U):
        f = E_U0 + i
        _f(ws, 'A%d' % f, nom)
        if p is None:
            _f(ws, 'B%d' % f, "=IF($B$%d=\"%s\",'%s'!$B$%d,'%s'!$B$%d)" % (E['modo'], MODOS[0], HO, S_T['c_comp'], HO, S_T['c_prop']), fmt=EUR3)
        else:
            _f(ws, 'B%d' % f, p, fmt=EUR3, fill=VERDE, locked=False)
        _f(ws, 'C%d' % f, 'ud')
        for j, (_n, (_pvp, _g, ud)) in enumerate(RECETAS.items()):
            _f(ws, '%s%d' % (PCOL[j], f), float(ud.get(nom, 0)), fmt='#,##0.0;-#,##0.0;"-"', fill=VERDE, locked=False)
        _f(ws, 'N%d' % f, origen, i=True)
    _f(ws, 'A%d' % E_R, 'RESULTADO POR PLATO (gramos y unidades en las columnas de arriba)', b=True, fill=GRIS)
    for j in range(NP):
        _f(ws, '%s%d' % (PCOL[j], E_R), '=%s%d' % (PCOL[j], E_CAB), b=True, fill=GRIS, wrap=True)
    ws.row_dimensions[E_R].height = 32
    kg = _kg_ext()
    for k, rot, fmt in RES:
        f = R[k]
        neg = k in ('c_tot', 'fc')
        _f(ws, 'A%d' % f, rot, b=neg)
        for j in range(NP):
            c = PCOL[j]
            rg = '%s$%d:%s$%d' % (c, E_G0, c, E_G1)
            ru = '%s$%d:%s$%d' % (c, E_U0, c, E_U1)
            ex = {
                'peso': 'SUM(%s)' % rg,
                'c_mat': 'SUMPRODUCT(%s,$B$%d:$B$%d)/%s+SUMPRODUCT(%s,$B$%d:$B$%d)' % (rg, E_G0, E_G1, kg, ru, E_U0, E_U1),
                'c_tot': '%s%d*(1+$B$%d)' % (c, R['c_mat'], E['merm']),
                'pvp_sin': '%s%d/(1+$B$%d)' % (c, E_PVP, E['iva']),
                'fc': '%s%d/%s%d' % (c, R['c_tot'], c, R['pvp_sin']),
                'sem': 'IF(%s%d<=$B$%d,"%s",IF(%s%d<=$B$%d+$B$%d,"%s","%s"))' % (
                    c, R['fc'], E['obj'], FC_OK, c, R['fc'], E['obj'], E['tol'], FC_VIG, FC_ALTO),
                'margen': '%s%d-%s%d' % (c, R['pvp_sin'], c, R['c_tot']),
                'pvp_sug': '%s%d/$B$%d*(1+$B$%d)' % (c, R['c_tot'], E['obj'], E['iva']),
            }[k]
            _f(ws, '%s%d' % (c, f), '=' + ex, fmt=fmt, b=neg, wrap=(k == 'sem'))
    ws.row_dimensions[R['sem']].height = 32
    rng = '%s%d:%s%d' % (PCOL[0], R['sem'], PCOL[-1], R['sem'])
    top = '%s%d' % (PCOL[0], R['sem'])
    for txt, fill in ((FC_OK, OK), (FC_VIG, AMBAR), (FC_ALTO, ROJO)):
        ws.conditional_formatting.add(rng, FormulaRule(formula=['%s="%s"' % (top, txt)], fill=fill))
    # margen por canal de la orden y la gringa
    _f(ws, 'A%d' % E_CH, 'MARGEN POR CANAL (orden y gringa)', b=True, fill=GRIS)
    for jj, col in CH_COL.items():
        _f(ws, '%s%d' % (col, E_CH), '=%s%d' % (PCOL[jj], E_CAB), b=True, fill=GRIS, wrap=True)
    ws.row_dimensions[E_CH].height = 32
    for k, rot, fmt in CANAL:
        f = CH[k]
        _f(ws, 'A%d' % f, rot, b=k.startswith('plat'))
        for jj, col in CH_COL.items():
            c = PCOL[jj]
            ps = '%s%d' % (c, R['pvp_sin'])
            ex = {
                'sala_e': '%s%d' % (c, R['margen']),
                'sala_p': '%s%d/%s' % (col, CH['sala_e'], ps),
                'llev_e': '%s%d-$B$%d' % (col, CH['sala_e'], E['env']),
                'llev_p': '%s%d/%s' % (col, CH['llev_e'], ps),
                'plat_e': '%s%d-$B$%d*%s' % (col, CH['llev_e'], E['com'], ps),
                'plat_p': '%s%d/%s' % (col, CH['plat_e'], ps),
            }[k]
            _f(ws, '%s%d' % (col, f), '=' + ex, fmt=fmt, b=k.startswith('plat'))
    _f(ws, 'A%d' % E_FIN, 'El food cost de la tabla es el del plato servido en sala. Para llevar se descuenta el envase del pedido; en plataforma, '
                          'además, la comisión sobre el precio sin IVA. Si en la plataforma el margen se queda corto, sube el precio allí o '
                          'vende la orden «para armar» (tortillas aparte), que aguanta mejor el viaje. Las bebidas no están en este escandallo.',
       i=True, wrap=True)
    ws.merge_cells('A%d:N%d' % (E_FIN, E_FIN))
    ws.row_dimensions[E_FIN].height = 48
    # validaciones
    dcol = lambda cols, fila: ['%s%d' % (c, fila) for c in cols]
    _dv(ws, 'decimal', ['B%d' % E[k] for k in ('iva', 'tol', 'merm', 'com')], f1='0', f2='0.9', op='between', titulo='Porcentaje no válido',
        msg='Escribe un porcentaje entre 0 y 0,9 (0,1 = 10 %).')
    _dv(ws, 'decimal', ['B%d' % E['obj']], f1='0.01', f2='0.9', op='between', titulo='Objetivo no válido',
        msg='El objetivo no puede ser 0 (se divide entre él): escribe entre 0,01 y 0,9.')
    _dv(ws, 'decimal', ['B%d' % E['env']] + ['B%d' % f for i, f in enumerate(range(E_G0, E_G1 + 1)) if not isinstance(ING_G[i][1], str)]
        + ['B%d' % (E_U0 + 1)] + dcol(PCOL, E_PVP)
        + ['%s%d' % (c, f) for c in PCOL for f in list(range(E_G0, E_G1 + 1)) + list(range(E_U0, E_U1 + 1))],
        f1='0', op='greaterThanOrEqual', titulo='Importe no válido', msg='Escribe un número mayor o igual que 0.')
    _dv(ws, 'list', ['B%d' % E['modo']], f1='"%s"' % ','.join(MODOS), titulo='Opción no válida', msg='Elige una opción de la lista.')
    _imprimir(ws, horizontal=True)


INSTRUCCIONES = [
    ('INSTRUCCIONES DE USO — Calculadora del trompo y escandallo por taco', True),
    ('Tres herramientas en un libro: la hoja «Trompo» te dice cuántos tacos salen de un cono al pastor y cuánto cuesta la carne de cada uno; '
     'la hoja «Tortilla» compara comprar la tortilla con hacerla tú y te da el equilibrio en tortillas al día; la hoja «Escandallo» calcula el '
     'coste, el food cost, el margen y el precio de carta de 8 tacos, una orden de 3 y una gringa, con el margen en sala, para llevar y en '
     'plataforma. Solo se escribe en las celdas VERDES; lo demás se calcula.', False),
    ('Todos los valores por defecto son una estimación del kit, editable: son ejemplos coherentes con el caso del plan de negocio, no precios '
     'de nadie. Cámbialos por los de tu proveedor y por lo que midas en tu cocina antes de decidir precios.', False),
    ('HOJA «TROMPO»', True),
    ('1. Cómo medir el cono. Pesa la carne cruda que montas en el trompo (antes del adobo) y apunta los kilos. Al cerrar, pesa lo que sobra del cono '
     'ya cocinado y anota a dónde va: se reaprovecha en un guiso o se tira. Con esos tres datos y los tacos que vendiste, la hoja te da el '
     'rendimiento real.', False),
    ('2. Cómo ajustar la merma. La merma de cocción y goteo de 35 % es una estimación del kit, sin fuente publicada. Mide tu primer cono: '
     'tacos servidos × gramos por taco ÷ 1.000 son los kilos servidos. Mueve la merma hasta que el rendimiento de la hoja coincida con el tuyo.', False),
    ('3. El resultado más útil es el coste por kg servido y el coste de carne por taco, que alimentan el escandallo. Compara tus tacos por kg '
     'crudo con la referencia de un puesto de la Ciudad de México, unos 25 por kg (Milenio). Si sirves más gramos, saldrán menos tacos.', False),
    ('4. «Carne cruda para los N tacos» te dice cuánta carne montar para un pedido grande. Las otras proteínas (carnitas, suadero, cochinita, '
     'barbacoa, tinga y campechano) usan su propio rendimiento tras la cocción, más bajo que el del cono.', False),
    ('5. Seguridad alimentaria. El cono montado y adobado es un preparado de carne: mantenlo a 4 °C como máximo hasta que entre al asador, mide '
     'con sonda y registra la temperatura en tu plan de APPCC. Pregunta a tu ayuntamiento y a tu comunidad por lo que exigen en tu zona.', False),
    ('HOJA «TORTILLA»', True),
    ('1. Escribe el precio del paquete, las piezas y el peso de cada tortilla, y los datos del nixtamal propio. Las tortillas por kg comprado '
     'son lo que necesitas para pedir al proveedor.', False),
    ('2. El equilibrio es el número de tortillas al día a partir del cual hacerlas compensa, contando a la persona que las hace y la inversión. '
     'Si gastas menos, compra. Con tu cifra de tortillas al día, la hoja te dice qué te conviene.', False),
    ('HOJA «ESCANDALLO»', True),
    ('1. Escribe el IVA (10 %), el food cost objetivo, el margen de vigilancia, la merma, el envase y la comisión de la plataforma, y elige si la '
     'tortilla de maíz es comprada o de nixtamal propio.', False),
    ('2. En la tabla, cada fila es un ingrediente con su precio por kilo sin IVA (o por unidad, en las tortillas). Los gramos de cada plato van '
     'en su columna; puedes cambiar los nombres, los gramos y los precios de carta. Las carnes toman su precio de la hoja «Trompo».', False),
    ('3. Para cada plato ves el coste con merma, el precio sin IVA, el food cost con semáforo, el margen y el precio de carta con IVA que '
     'necesitarías para llegar a tu objetivo. Abajo, el margen de la orden y de la gringa en sala, para llevar y en plataforma.', False),
    ('PRECIOS, IVA Y FUENTES', True),
    ('El IVA de la comida preparada para consumir en el local, para llevar o a domicilio es el 10 % (Ley 37/1992, art. 91.Uno.1.1.º y 2.2.º). '
     'Las bebidas alcohólicas y las aguas frescas o refrescos con azúcares o edulcorantes añadidos que salen del local llevan el 21 % y no '
     'están en este escandallo.', False),
    ('Los precios y rendimientos salen del estudio de mercado del kit (octubre 2026) y de estimaciones del kit; la referencia de tacos por kg '
     'viene de Milenio. Ninguna cifra de esta calculadora es un dato publicado de tu negocio: sustitúyelas por las tuyas.', False),
]


def hoja_instrucciones(wb):
    ws = wb.create_sheet(HI)
    ws.column_dimensions['A'].width = 110
    f = 1
    for t, b in INSTRUCCIONES:
        _f(ws, 'A%d' % f, t, b=b, sz=13 if f == 1 else 11, wrap=True)
        if not b:
            ws.row_dimensions[f].height = max(16, 15 * math.ceil(len(t) / 120.0))
        f += 2 if b and f > 1 else 1
    f += 1
    for t in ('Para editar una celda que no esté en verde: Revisar → Desproteger hoja (no tiene contraseña).',
              'Más plantillas y kits del catálogo en aichef.pro/productos-digitales', motor.BIO_LINE,
              motor.VERSION_FMT.format(pid=PID)):
        _f(ws, 'A%d' % f, t, wrap=True)
        f += 1
    _imprimir(ws)


def generar(destino):
    wb = openpyxl.Workbook()
    hoja_trompo(wb)
    hoja_tortilla(wb)
    hoja_escandallo(wb)
    hoja_instrucciones(wb)
    pr = wb.properties
    pr.title = '%s · v%s' % (TITULO, motor.VERSION)
    pr.subject = 'Plan de Negocio: Taquería'
    pr.description = 'aichef.pro/%s' % PID
    pr.keywords = 'plan negocio taquería, calculadora trompo al pastor, escandallo taco, AI Chef Pro'
    pr.creator = pr.lastModifiedBy = motor.CREATOR
    os.makedirs(os.path.dirname(destino), exist_ok=True)
    wb.save(destino)
    sys.path.insert(0, PD)
    import inject_cache as IC                                                      # noqa: E402
    nf, inj, fail = IC.inject(destino)
    print('%s: formulas=%d cache_inyectado=%d fallos_pycel=%d' % (os.path.basename(destino), nf, inj, fail))
    return fail


def valores_expuestos(ruta):
    """Lee de la CACHÉ del xlsx los 5 valores expuestos: {nombre: valor}."""
    wv = openpyxl.load_workbook(ruta, data_only=True)
    return OrderedDict((n, wv[h][c].value) for n, (h, c) in EXPUESTOS.items())


# ---------------------------------------------------------------------------------------------------------------------
# Autotest
# ---------------------------------------------------------------------------------------------------------------------
PROHIBIDOS = [r'pizz', r'napol[ei]t', r'margherita', r'\ble[ñn]a\b', r'levadura', r'\bbolas?\b', r'sin gluten', r'fior di latte',
              r'mozzarella', r'\bETG\b', r'pliego', r'\bfreelance', r'validad[oa]s?', r'carnet', r'certificad', r'service charge|sales tax']
CELDA = {}                                                                         # clave de entrada → (hoja, celda)
for _k, _f2 in M.items():
    CELDA[_k] = (HT, 'B%d' % _f2)
for _i in range(NB):
    for _c, _k in zip('BCDE', ('bk', 'bp', 'br', 'bg')):
        CELDA['%s%d' % (_k, _i)] = (HT, '%s%d' % (_c, B0 + _i))
for _k, _f2 in T.items():
    CELDA[_k] = (HO, 'B%d' % _f2)
for _k, _f2 in E.items():
    CELDA[_k] = (HE, 'B%d' % _f2)
for _i in range(NG):
    CELDA['p%d' % _i] = (HE, 'B%d' % (E_G0 + _i))
CELDA['th'] = (HE, 'B%d' % (E_U0 + 1))
for _j in range(NP):
    CELDA['pvp%d' % _j] = (HE, '%s%d' % (PCOL[_j], E_PVP))
    for _i in range(NG):
        CELDA['g%d_%d' % (_i, _j)] = (HE, '%s%d' % (PCOL[_j], E_G0 + _i))
    for _i in range(NU):
        CELDA['u%d_%d' % (_i, _j)] = (HE, '%s%d' % (PCOL[_j], E_U0 + _i))


def celdas_salida(r):
    """(hoja, celda, valor esperado, etiqueta) de todas las salidas."""
    out = []
    for k in S_A:
        out.append((HT, 'B%d' % S_A[k], r[k], k))
    for i, b in enumerate(r['B']):
        for col, k in zip('FGHI', ('coc', 'ckg', 'tacos', 'ctaco')):
            out.append((HT, '%s%d' % (col, B0 + i), b[k], 'guiso%d.%s' % (i, k)))
    for k in S_T:
        out.append((HO, 'B%d' % S_T[k], r['T'][k], 'T.' + k))
    for i in range(NG):
        out.append((HE, 'B%d' % (E_G0 + i), r['pg'][i], 'precio%d' % i))
    for i in range(NU):
        out.append((HE, 'B%d' % (E_U0 + i), r['pu'][i], 'precio_ud%d' % i))
    for j in range(NP):
        for k in R:
            out.append((HE, '%s%d' % (PCOL[j], R[k]), r['p%d' % j][k], '%s[%d]' % (k, j)))
    for jj, col in CH_COL.items():
        for k in CH:
            out.append((HE, '%s%d' % (col, CH[k]), r['p%d' % jj][k], '%s[%d]' % (k, jj)))
    return out


def _distinto(x, v):
    if isinstance(v, str):
        return x != v
    return not isinstance(x, (int, float)) or abs(float(x) - v) > 0.01


def _checks_logica(nombre, jj, err):
    r = calcular(jj)
    if abs(r['servidos'] - r['rend'] * jj['K']) > 1e-6:
        err.append('%s: servidos ≠ rendimiento × K' % nombre)
    if abs(r['tacos'] * jj['gt'] - r['servidos'] * jj['KG']) > 0.5:
        err.append('%s: tacos × g ≠ servidos × 1.000' % nombre)
    tpd = r['T']['eq']
    if tpd > 0:
        tt = calcular(dict(jj, tpd=tpd))['T']
        if abs(tt['dia_comp'] - tt['dia_prop']) > 0.01:
            err.append('%s: en el equilibrio, coste diario comprada %.4f ≠ propia %.4f' % (nombre, tt['dia_comp'], tt['dia_prop']))
    for j in (J_ORDEN, J_GRINGA):
        p = r['p%d' % j]
        if not (p['plat_e'] < p['llev_e'] < p['sala_e'] and p['plat_p'] < p['llev_p'] < p['sala_p']):
            err.append('%s: margen plataforma < llevar < sala no se cumple en la columna %d' % (nombre, j))


def autotest(destino):
    from pycel import ExcelCompiler
    err, n = [], 0
    wv = openpyxl.load_workbook(destino, data_only=True)
    wf = openpyxl.load_workbook(destino)
    base = valores_defecto()
    j2 = dict(base, K=20.0, a=0.10, Pc=6.20, Pa=7.0, pk=1.5, pp=2.4, mc=0.30, mr=0.07, S=1.2, r=0.5, gt=30.0, N=250,
              bk0=12.0, bp1=8.2, br2=0.58, bg3=45.0, tp=11.0, tn=80, tg=25.0, mz=1.0, mm=2.2, cg=1.5, gn=25.0, tw=15000.0, ti=9000.0, ty=5,
              td=280, tpd=900, iva=0.10, obj=0.28, tol=0.03, merm=0.08, modo=MODOS[1], env=0.7, com=0.25, th=0.25, p7=1.5, p14=9.0,
              pvp0=3.2, pvp8=9.0, g0_1=40.0, g3_5=7.0, g1_9=20.0)
    j3 = dict(base, K=8.0, a=0.06, Pc=5.0, mc=0.40, mr=0.04, gt=40.0, N=30, bk3=6.0, br4=0.7, tp=20.0, tn=100, tg=20.0, mz=1.4, mm=1.8,
              cg=2.5, gn=20.0, tw=21000.0, ti=25000.0, ty=8, td=320, tpd=150, obj=0.2, tol=0.1, merm=0.0, env=0.4, com=0.35, th=0.18,
              p9=2.5, pvp9=8.0, g0_8=120.0, g2_2=40.0)
    # 1. caché y pycel del juego por defecto; pycel de los otros dos con set_value
    xl = ExcelCompiler(filename=destino)
    r0 = calcular(base)
    for hoja, c, v, lab in celdas_salida(r0):
        n += 1
        for et, x in (('pycel', xl.evaluate("'%s'!%s" % (hoja, c))), ('caché', wv[hoja][c].value)):
            if _distinto(x, v):
                err.append('defecto · %s!%s (%s): %s %r ≠ esperado %r' % (hoja, c, lab, et, x, v))
    for nombre, jj in (('juego 2', j2), ('juego 3', j3)):
        r = calcular(jj)
        xl2 = ExcelCompiler(filename=destino)
        for hoja, c, v, lab in celdas_salida(r0):
            xl2.evaluate("'%s'!%s" % (hoja, c))
        for k, v in jj.items():
            if v != base[k]:
                h, c = CELDA[k]
                xl2.set_value("'%s'!%s" % (h, c), v)
        for hoja, c, v, lab in celdas_salida(r):
            n += 1
            x = xl2.evaluate("'%s'!%s" % (hoja, c))
            if _distinto(x, v):
                err.append('%s · %s!%s (%s): pycel %r ≠ esperado %r' % (nombre, hoja, c, lab, x, v))
    # 2. identidades de la lógica en los tres juegos (y servidos / tacos también sobre la hoja por defecto)
    for nombre, jj in (('defecto', base), ('juego 2', j2), ('juego 3', j3)):
        n += 3
        _checks_logica(nombre, jj, err)
    kg = xl.evaluate("'%s'!B%d" % (HT, M['KG']))
    k_ = xl.evaluate("'%s'!B%d" % (HT, M['K']))
    sv, rd = xl.evaluate("'%s'!B%d" % (HT, S_A['servidos'])), xl.evaluate("'%s'!B%d" % (HT, S_A['rend']))
    tc, gt = xl.evaluate("'%s'!B%d" % (HT, S_A['tacos'])), xl.evaluate("'%s'!B%d" % (HT, M['gt']))
    n += 2
    if abs(sv - rd * k_) > 1e-6 or abs(tc * gt - sv * kg) > 0.5:
        err.append('hoja: servidos ≠ rendimiento × K o tacos × g ≠ servidos × 1.000')
    # 3. valores con sentido
    pastor = r0['p%d' % J_PASTOR]
    if not 0.18 <= pastor['fc'] <= 0.30:
        err.append('food cost del taco al pastor %.1f %% fuera de 18-30 %%' % (pastor['fc'] * 100))
    if abs(r0['rend'] - 0.648) > 0.001 or not 277 <= r0['tacos'] <= 279 or abs(r0['tacos_kg'] - 18.5) > 0.1:
        err.append('rendimiento %.4f / tacos %.1f / tacos por kg %.2f fuera de lo esperado (64,8 %% · 278 · 18,5)'
                   % (r0['rend'], r0['tacos'], r0['tacos_kg']))
    if abs(r0['coste_kg'] - 9.23) > 0.01 or abs(r0['carne_taco'] - 0.32) > 0.005:
        err.append('coste/kg servido %.3f o carne por taco %.4f fuera de 9,23 / 0,32' % (r0['coste_kg'], r0['carne_taco']))
    if not 0.60 <= pastor['c_tot'] <= 0.62:
        err.append('coste del taco al pastor %.4f fuera de 0,60-0,62' % pastor['c_tot'])
    if not 650 <= r0['T']['eq'] <= 750 or r0['T']['sem'] != SEM_COMPRA:
        err.append('equilibrio %.0f fuera de 650-750 o semáforo ≠ compra con 280 al día' % r0['T']['eq'])
    if calcular(dict(base, tpd=900))['T']['sem'] != SEM_PROPIA:
        err.append('el semáforo de la tortilla no pasa a «Hazla tú» por encima del equilibrio')
    # 4. los 5 valores expuestos: caché frente a calcular()
    esperados = {'TACOS_POR_CONO': r0['tacos'], 'RENDIMIENTO_TROMPO': r0['rend'], 'COSTE_TACO_PASTOR': pastor['c_tot'],
                 'FC_TACO_PASTOR': pastor['fc'], 'EQUILIBRIO_TORTILLAS_DIA': r0['T']['eq']}
    expuestos = valores_expuestos(destino)
    for nombre, e_ in esperados.items():
        n += 1
        if _distinto(expuestos[nombre], e_):
            err.append('expuesto %s: caché %r ≠ %r' % (nombre, expuestos[nombre], e_))
    # 5. formato, protección, DV, fórmulas, textos
    if motor.gate_formatos(wf, FICHERO):
        err.append('gate_formatos: %r' % motor.gate_formatos(wf, FICHERO)[:3])
    dvh = dict((w.title, dict((str(rng), dv) for dv in w.data_validations.dataValidation for rng in dv.sqref.ranges))
               for w in (wf[HT], wf[HO], wf[HE]))
    dvs = dict(dvh[HT], **dvh[HE])
    for hoja, k in ((HT, 'K'), (HT, 'N'), (HT, 'mc'), (HO, 'tpd'), (HO, 'td'), (HO, 'tp'), (HE, 'obj'), (HE, 'modo'), (HE, 'iva')):
        c = CELDA[k][1]
        if c not in dvh[hoja]:
            err.append('DV ausente en %s!%s' % (hoja, c))
    dv = dvh[HT].get(CELDA['N'][1])
    if dv is None or dv.type != 'whole':
        err.append('DV del número de tacos (entero) ausente o incorrecta')
    dv = dvh[HE].get(CELDA['modo'][1])
    if dv is None or dv.type != 'list':
        err.append('DV de lista del modo de tortilla ausente')
    textos, todo_celdas = [], []
    for w in wf.worksheets:
        if w.page_setup.paperSize != 9 or not w.protection.sheet:
            err.append('%s: A4 y protección' % w.title)
        for row in w.iter_rows():
            for c in row:
                v = c.value
                if v is not None:
                    todo_celdas.append(str(v))
                if isinstance(v, str) and c.data_type != 'f':
                    textos.append(v)
                if c.data_type == 'f':
                    if wv[w.title][c.coordinate].value is None:
                        err.append('%s!%s: fórmula sin caché' % (w.title, c.coordinate))
                    sin_refs = re.sub(r'"[^"]*"', '', re.sub(r"'[^']+'!|\$?[A-Z]{1,2}\$?\d+", '', str(c.value)))
                    for num in re.findall(r'(?<![A-Za-z])\d+\.?\d*', sin_refs):
                        if num not in ('0', '1'):
                            err.append('%s!%s: número tecleado %s en %s' % (w.title, c.coordinate, num, c.value))
                if c.fill.fgColor.rgb in ('00E8F5E9', 'FFE8F5E9') and c.protection.locked:
                    err.append('%s!%s: verde bloqueada' % (w.title, c.coordinate))
                if c.fill.fgColor.rgb in ('00E8F5E9', 'FFE8F5E9') and c.data_type == 'f':
                    err.append('%s!%s: verde con fórmula' % (w.title, c.coordinate))
    todo = '\n'.join(textos)
    barrido = '\n'.join(todo_celdas)
    for rx in PROHIBIDOS:
        if re.search(rx, barrido, re.I):
            err.append('prohibido /%s/' % rx)
    if re.search(r'[^\x00-ɏ‐-‧€−→·]', barrido):
        err.append('caracteres no latinos o raros: %r' % sorted(set(re.findall(r'[^\x00-ɏ‐-‧€−→·]', barrido))))
    if 'Estimación del kit, editable' not in barrido and 'estimación del kit, editable' not in barrido:
        err.append('falta el rótulo «estimación del kit, editable»')
    if motor.VERSION_FMT.format(pid=PID) not in todo or motor.BIO_LINE not in todo:
        err.append('falta la línea de versión o la bio')
    if '1.0' not in motor.VERSION_FMT or 'noviembre 2026' not in motor.VERSION_FMT:
        err.append('versión no es 1.0 · noviembre 2026: %s' % motor.VERSION_FMT)
    if wf.sheetnames != [HT, HO, HE, HI]:
        err.append('pestañas %r' % wf.sheetnames)
    if NG + NU < 16:
        err.append('menos de 16 ingredientes')
    p = wf.properties
    if p.creator != 'AI Chef Pro' or p.description != 'aichef.pro/%s' % PID or 'AI Chef Pro' not in (p.keywords or ''):
        err.append('docProps: creator=%r description=%r keywords=%r' % (p.creator, p.description, p.keywords))
    for x in err:
        print('ERROR', x)
    print('autotest calculadora taquería ES: %d comprobaciones (pycel + caché por defecto; pycel en 2 juegos más; identidades) · %s'
          % (n, 'VERDE' if not err else 'ROJO (%d errores)' % len(err)))
    print('  cono: %.0f kg → %.2f kg servidos · rendimiento %.1f %% · %.1f tacos (%.1f por kg crudo) · %.3f €/kg servido · carne %.3f €/taco'
          % (base['K'], r0['servidos'], r0['rend'] * 100, r0['tacos'], r0['tacos_kg'], r0['coste_kg'], r0['carne_taco']))
    for nombre, (h, c) in EXPUESTOS.items():
        print('  %-26s %s!%-5s %s' % (nombre, h, c, expuestos[nombre]))
    for j, nom in enumerate(RECETAS):
        q = r0['p%d' % j]
        print('  %-22s coste %.3f € · food cost %5.1f %% · %s' % (nom, q['c_tot'], q['fc'] * 100, q['sem']))
    for j in (J_ORDEN, J_GRINGA):
        q = r0['p%d' % j]
        print('  canal %-20s sala %.2f € (%.0f %%) · llevar %.2f € (%.0f %%) · plataforma %.2f € (%.0f %%)'
              % (list(RECETAS)[j], q['sala_e'], q['sala_p'] * 100, q['llev_e'], q['llev_p'] * 100, q['plat_e'], q['plat_p'] * 100))
    return 1 if err else 0


if __name__ == '__main__':
    ap = argparse.ArgumentParser()
    ap.add_argument('--salida', required=True, help='carpeta de salida')
    ap.add_argument('--autotest', action='store_true')
    a = ap.parse_args()
    destino = os.path.join(a.salida, FICHERO)
    if not a.autotest and generar(destino):
        sys.exit('ERROR: fallos de pycel en inject_cache')
    sys.exit(autotest(destino))
