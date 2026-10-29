#!/usr/bin/env python3
"""
semilla_taq.py — Plan de Negocio: Taquería (ES 1.0; F1-camino §2-§3): prepara el ORIGEN que `planes-v2_0/main.py`
convierte en el plan financiero y el checklist de la taquería. Copia adaptada de `plan-pizzeria/semilla_piz.py`.

Copia los dos xlsx de la Pizzería ES 1.0 publicados (`dl/plan-negocio-pizzeria/`, SOLO plan y checklist, SOLO LECTURA;
molde A-β + checklist C2) con los nombres de fichero de la taquería y escribe lo que el módulo de contenido no alcanza
(F1-camino §1, medido el 29-oct):
  · plan: valores de `0. Supuestos` (grupo_a conserva lo verde que trae el libro; incluye `B60`/`B61` = 7/12, las
    pagas con el Mes 1 en enero: la Pizzería trae 11/4), títulos de hoja (Inversión A1/A2, PyG A2, Punto Equilibrio
    A2), las 18 partidas de inversión (filas 7-24: rótulo, importe, nota, ¿IVA?), los 9 costes fijos propios del P&L
    (filas 26-34: rótulo e importe; la nota la pone el contenido) y los docProps description/keywords. `PyG!G13` ya es
    VALOR en la Pizzería y lo rehace el motor con `IVA_COMPRAS`: aquí NO se toca (el assert de fórmula de la Pizzería
    rompería);
  · checklist: título, pestañas, cabeceras de fase y las tareas de `datos_checklist.py` con el pie «Tareas
    completadas:» debajo; `Instrucciones!A3` con el nombre del producto; DV, formato condicional y contadores los rehace
    el motor.
Corre en el VPS (openpyxl). Después: `main.py --producto plan-negocio-taqueria --dry-run --origen <salida>` (ciclo.sh).

    /root/venv-guias/bin/python semilla_taq.py --salida /root/tmp-taq-f2/motor/origen

Sesión Claude Code, 29-oct-2026.
"""
import argparse
import copy
import importlib
import os
import shutil
import sys

import openpyxl

import datos_checklist as DC

AQUI = os.path.dirname(os.path.abspath(__file__))
REPO = os.path.abspath(os.path.join(AQUI, '..', '..', '..'))
PIZ = os.path.join(REPO, 'astro-site', 'public', 'dl', 'plan-negocio-pizzeria')
PLAN_PIZ, CHECK_PIZ = 'plan-financiero-pizzeria.xlsx', 'checklist-apertura-pizzeria.xlsx'
PLAN, CHECK = 'plan-financiero-taqueria.xlsx', 'checklist-apertura-taqueria.xlsx'
PID = 'plan-negocio-taqueria'
NOMBRE = 'Plan de Negocio: Taquería'

TITULOS_PLAN = {
    ('Inversión Inicial', 'A1'): 'PLAN DE NEGOCIO: TAQUERÍA',
    ('Inversión Inicial', 'A2'): 'Inversión inicial estimada — taquería con trompo, barra de salsas, 32 plazas y terraza',
    ('Inversión Inicial', 'A3'): 'Ajusta las cifras a tu proyecto: son las de un caso tipo en una ciudad media, sin '
                                 'traspaso. Pon tus presupuestos reales en las celdas de importe.',
    ('PyG 3 Años', 'A2'): 'Escenario realista — taquería con trompo, barra de salsas y para llevar',
    ('Punto Equilibrio', 'A2'): 'Punto de equilibrio mensual y diario — taquería con trompo',
}

#: Filas 7-24 de 'Inversión Inicial': (rótulo, importe, nota, ¿lleva IVA?). R §3: 19 partidas en 18 filas («Rótulo y
#: fachada», 2.500 €, se funde con la obra: F1-camino §2), 95.500 € (obra 48.500 · maquinaria 36.000 · no amortizable
#: 11.000) + imprevistos del 10 % sobre la obra (fila 27, por fórmula) = 100.350 €. Sin palabras de recuento en el
#: rótulo (motor.RX_RECUENTO), sin paréntesis con número. ⚠️ «campana» a secas la acentúa `motor.TILDES` a «campaña»:
#: siempre «campana extractora»; y nunca el verbo «publica» (→ «pública»).
INVERSION = [
    ('Proyecto técnico de actividad y de salida de humos, y dirección de obra', 3500,
     'Proyecto de la actividad y del conducto de humos, memoria para la licencia o la declaración responsable y '
     'certificado final de obra', 'Sí'),
    ('Obra y acondicionamiento del local, rótulo y fachada', 26500,
     'Cocina alicatada, suelo antideslizante, aseos (uno adaptado), falsos techos y aislamiento acústico, más el rótulo '
     'y la fachada (unos 2.500 €). Estimación del kit: pide dos presupuestos', 'Sí'),
    ('Conducto de extracción exclusivo hasta cubierta', 6000,
     'Conducto exclusivo hasta la cubierta para el trompo, la plancha, el comal y la freidora: Madrid pide superar en '
     '1 m lo que haya en 15 m (Ordenanza 4/2021). Con trompo o parrilla de carbón, además, filtro de partículas. '
     'Ningún fabricante da precio por metro: pide oferta con tu proyecto', 'Sí'),
    ('Instalación eléctrica, cuadro y boletín', 4500,
     'Frío, extracción, climatización, lavavajillas y cafetera. Pregunta la potencia disponible ANTES de firmar: una '
     'acometida nueva tarda', 'Sí'),
    ('Fontanería, separador de grasas, climatización e instalación de gas', 8000,
     'Agua fría y caliente en la cocina, desagües con separador de grasas, climatización de la sala y la instalación '
     'de gas del trompo, la plancha, el comal y los fogones por empresa instaladora habilitada, con su certificado '
     '(RD 919/2006)', 'Sí'),
    ('Trompo: asador vertical de gas de cuatro quemadores', 1800,
     'Para conos de 30-35 kg, con bandeja, cuchillos de corte y base para la piña. En el catálogo de un distribuidor '
     'de hostelería, de 1.279 a 1.784 € sin IVA (octubre de 2026)', 'Sí'),
    ('Plancha de acero a gas y comal', 3000,
     'Plancha lisa de acero para las carnes y comal para calentar la tortilla. Una plancha a gas de 40 cm ronda los '
     '1.040 €; la de 80 cm a gas no tiene precio de referencia: estimación del kit, pide oferta', 'Sí'),
    ('Freidora, fogón de carnitas, baño maría y mantenedor de tortillas', 3700,
     'Freidora para los totopos y los tacos dorados, fogón para las carnitas y los guisados, baño maría GN y '
     'mantenedor de tortillas. Estimación del kit', 'Sí'),
    ('Campana extractora con motor y filtros', 1500,
     'Campana extractora de 2,5 m con motor y filtros de grasa sobre el trompo, la plancha, el comal y la freidora '
     '(una de 1,5-2 m ronda los 815 € sin IVA)', 'Sí'),
    ('Frío: armarios, congelador y mesa refrigerada de salsas', 4500,
     'Armarios refrigerados (unos 900 € sin IVA el de 400 l), congelador y mesa refrigerada de granito para la barra '
     'de salsas y las guarniciones (de 1.180 a 1.950 € sin IVA)', 'Sí'),
    ('Lavavajillas de capota, fregaderos, lavamanos, mesas inox y estanterías', 5500,
     'Lavavajillas, fregaderos, lavamanos de accionamiento no manual, mesas de acero inoxidable y estanterías', 'Sí'),
    ('Utillaje y menaje: licuadora, molcajetes, sondas y prensa de tortillas', 3500,
     'Licuadora industrial para las salsas, molcajetes, termómetros sonda, básculas, recipientes GN, vajilla, cestas, '
     'vitroleros para las aguas frescas, una prensa de tortillas (unos 48 € con IVA) y la cartelería de alérgenos', 'Sí'),
    ('Barra: botellero frío, tirador de cerveza y cafetera', 4000,
     'Botellero frío, tirador de cerveza y una cafetera pequeña. Si la cervecera cede el equipo de tiraje, esta '
     'partida baja', 'Sí'),
    ('Mobiliario de sala, barra y terraza', 6500,
     'Mesas, sillas y taburetes para 24 plazas en mesa y 8 en barra, y 16 en la terraza de temporada', 'Sí'),
    ('TPV con Veri*factu, datáfonos, impresora de comandas y tablet', 2000,
     'Software de facturación conforme a Veri*factu, datáfonos, impresora de cocina y la tablet de las plataformas',
     'Sí'),
    ('Constitución de la SL, licencia, tasas, registro sanitario y APPCC', 3000,
     'SL unipersonal, licencia o declaración responsable, tasas municipales, registro sanitario autonómico y '
     'redacción del plan APPCC. Las tasas no llevan IVA; los honorarios sí: cambia la columna de al lado si tu '
     'presupuesto es sobre todo de honorarios', 'No'),
    ('Marca, web, fotos y lanzamiento', 4500,
     'Identidad, web con la carta, sesión de fotos y campaña de apertura', 'Sí'),
    ('Stock inicial: carnes, chiles, tortillas, bebida y envases', 3500,
     'Carnes, chiles secos y producto mexicano importado (pídelo con su plazo de entrega: tarda semanas), tortillas, '
     'bebida, mezcal y envases para llevar', 'Sí'),
]
assert len(INVERSION) == 18 and sum(x[1] for x in INVERSION) == 95500, sum(x[1] for x in INVERSION)

#: Filas 26-34 de 'PyG 3 Años' (costes fijos propios): 9 filas, una por fila del molde (no se suprimen ni se añaden:
#: desplazaría el P&L), 13.800 €/año (R §3: 10 filas y 12.600 €; la formación entra en la PRL y la tasa de terraza de
#: 1.200 € en la fila de residuos: F1-camino §1, vuelta 1). Sin palabras de recuento en el rótulo. Las claves de
#: `contenido….FIJOS` son estos rótulos normalizados.
FIJOS = {
    26: ('Gestoría laboral, fiscal y contable', 2400),
    27: ('Mantenimiento de la campana extractora, el gas y el frío', 1800),
    28: ('Marketing y publicidad', 3500),
    29: ('Software, TPV, teléfono e internet', 1500),
    30: ('Limpieza, lavandería, EPI y análisis de autocontrol', 1200),
    31: ('Tasa de terraza y gestión de residuos (pon las de tu ayuntamiento)', 1700),
    32: ('Desinsectación, desratización y desinfección (DDD)', 450),
    33: ('Derechos de autor por música ambiental (SGAE y AGEDI-AIE)', 500),
    34: ('Prevención de riesgos laborales, vigilancia de la salud y formación', 750),
}
assert sum(x[1] for x in FIJOS.values()) == 13800


def _contenido():
    sys.path.insert(0, os.path.join(os.path.dirname(AQUI), 'planes-v2_0'))
    import grupo_a
    return grupo_a, importlib.import_module('contenido_plan_negocio_taqueria.a')


def _supuestos(wb):
    """Los valores de `0. Supuestos` del caso taquería: `grupo_a` CONSERVA lo que el libro ya trae en las celdas
    verdes y sólo usa el valor del contenido en una celda vacía; sin este paso el caso seguiría siendo el de la
    pizzería (66 × 19,50 € × 300, pagas 11/4)."""
    grupo_a, contenido = _contenido()
    coords = dict((t[0], t[1]) for t in grupo_a.SUPUESTOS_BASE)
    ws = wb['0. Supuestos']
    escritas = 0
    for clave, t in contenido.SUPUESTOS.items():
        if t[2] is None:
            continue
        assert clave in coords, clave
        ws[coords[clave]] = t[2]
        escritas += 1
    # F1-camino §2 y §7 F2 (2): las pagas con el Mes 1 en enero (la Pizzería trae 11/4, Mes 1 = septiembre)
    assert coords['mes_paga_1'] == 'B60' and coords['mes_paga_2'] == 'B61', (coords['mes_paga_1'], coords['mes_paga_2'])
    assert (ws['B60'].value, ws['B61'].value) == (7, 12), (ws['B60'].value, ws['B61'].value)
    return escritas


def _docprops(wb):
    """El motor no reescribe `description` ni `keywords` y heredaba los de la pizzería."""
    wb.properties.description = 'aichef.pro/' + PID
    wb.properties.keywords = 'plan de negocio taquería, taqueria, trompo al pastor, AI Chef Pro'


def plan(salida):
    destino = os.path.join(salida, PLAN)
    shutil.copy2(os.path.join(PIZ, PLAN_PIZ), destino)
    wb = openpyxl.load_workbook(destino)
    print('  supuestos escritos:', _supuestos(wb))
    _docprops(wb)
    for (hoja, coord), texto in TITULOS_PLAN.items():
        wb[hoja][coord] = texto
    ws = wb['Inversión Inicial']
    for i, (rot, imp, nota, iva) in enumerate(INVERSION):
        r = 7 + i
        assert not str(ws.cell(r, 1).value).lower().startswith(('fianza', 'alquiler', 'imprevistos')), r
        ws.cell(r, 1).value, ws.cell(r, 2).value, ws.cell(r, 4).value, ws.cell(r, 5).value = rot, imp, nota, iva
    assert str(ws['A25'].value).lower().startswith('fianza'), ws['A25'].value
    assert str(ws['A26'].value).lower().startswith('alquiler'), ws['A26'].value
    assert str(ws['A27'].value).lower().startswith('imprevistos'), ws['A27'].value
    ws = wb['PyG 3 Años']
    assert str(ws['A25'].value).startswith('Seguros') and str(ws['A35'].value).startswith('Amortización'), (
        ws['A25'].value, ws['A35'].value)
    for r, (rot, imp) in FIJOS.items():
        ws.cell(r, 1).value, ws.cell(r, 2).value = rot, imp
        ws.cell(r, 6).value = None          # la nota la pone el contenido (FIJOS)
    # `PyG!G13` (IVA soportado de la compra de comida): en la Pizzería ya es VALOR (0,064) y el motor lo reescribe con
    # `IVA_COMPRAS` (0,085). No se toca aquí: el assert de fórmula de `semilla_piz.py` rompería (F1-camino §3).
    assert str(ws['A13'].value).startswith('Coste de mercancía — comida'), ws['A13'].value
    assert not str(ws['G13'].value).startswith('='), ws['G13'].value
    wb.save(destino)
    return destino


def _copiar_estilo(origen, destino):
    destino.font, destino.border = copy.copy(origen.font), copy.copy(origen.border)
    destino.fill, destino.alignment = copy.copy(origen.fill), copy.copy(origen.alignment)
    destino.number_format, destino.protection = origen.number_format, copy.copy(origen.protection)


def checklist(salida):
    destino = os.path.join(salida, CHECK)
    shutil.copy2(os.path.join(PIZ, CHECK_PIZ), destino)
    wb = openpyxl.load_workbook(destino)
    fases = [ws for ws in wb.worksheets if ws.title[:2] in ('F1', 'F2', 'F3', 'F4', 'F5', 'F6')]
    assert len(fases) == 6
    for ws, (titulo, cabecera, tareas) in zip(fases, DC.FASES):
        ws.title = titulo
        ws['A1'], ws['A2'] = DC.TITULO, cabecera
        assert ws['B4'].value and 'TRÁMITE' in ws['B4'].value.upper(), ws['B4'].value
        modelo = [copy.copy(ws.cell(5, c)) for c in range(1, 6)]
        alto = ws.row_dimensions[5].height
        pie = [copy.copy(ws.cell(r, c)) for r in range(5, ws.max_row + 1) for c in range(1, 6)
               if isinstance(ws.cell(r, c).value, str) and 'completad' in ws.cell(r, c).value.lower()]
        assert pie, ws.title
        fila_pie = pie[0].row
        estilo_pie = [copy.copy(ws.cell(fila_pie, c)) for c in range(1, 6)]
        # vaciar la zona de tareas y el pie (valores, DV y formato condicional los rehace el motor)
        for r in range(5, ws.max_row + 1):
            for c in range(1, 6):
                ws.cell(r, c).value = None
            ws.row_dimensions[r].height = None
        ws.data_validations.dataValidation = []
        ws.conditional_formatting = type(ws.conditional_formatting)()
        for i, (tarea, resp, plazo, nota) in enumerate(tareas):
            r = 5 + i
            for c, v in enumerate(('☐', tarea, resp, plazo, nota), start=1):
                cel = ws.cell(r, c)
                _copiar_estilo(modelo[c - 1], cel)
                cel.value = v
            ws.row_dimensions[r].height = alto
        r = 5 + len(tareas) + 1
        for c in range(1, 6):
            _copiar_estilo(estilo_pie[c - 1], ws.cell(r, c))
        ws.cell(r, 2).value = 'Tareas completadas:'
        ws.cell(r, 3).value = '=COUNTIF(A5:A%d,"✓")' % (4 + len(tareas))
        ws.cell(r, 4).value = 'de'
        ws.cell(r, 5).value = '=COUNTIF(B5:B%d,"?*")-COUNTIF(A5:A%d,"N/A")' % (4 + len(tareas), 4 + len(tareas))
    # `motor.asegurar_instrucciones` no reescribe la hoja si ya existe: la de la Pizzería nombraba su producto
    ins = wb['Instrucciones']
    assert 'Pizzería Napolitana' in str(ins['A3'].value), ins['A3'].value
    ins['A3'] = 'Este fichero forma parte del «%s» de AI Chef Pro.' % NOMBRE
    _docprops(wb)
    wb.save(destino)
    return destino


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--salida', required=True)
    a = ap.parse_args()
    if os.path.isdir(a.salida):
        shutil.rmtree(a.salida)
    os.makedirs(a.salida)
    print('plan      →', plan(a.salida))
    print('checklist →', checklist(a.salida), '(%d tareas: %s)' % (
        sum(len(t) for _, _, t in DC.FASES), '/'.join(str(len(t)) for _, _, t in DC.FASES)))


if __name__ == '__main__':
    main()
