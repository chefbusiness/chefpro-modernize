#!/usr/bin/env python3
"""
extraer_textos.py — Restaurant Inventory Kit Pro (EN) · cimientos de la F2-EN.

Copia adaptada del extractor del piloto (`recipe-costing-kit/extraer_textos.py`). Lee los 9 xlsx PUBLICADOS
del Kit de Inventario v2.0 (`astro-site/public/dl/kit-inventario/`, SOLO LECTURA) y escribe, junto a este
script (o en --salida):

  · textos_es.json  cadenas de texto únicas con su contexto (fichero, hoja, celda o «dónde»: valor de
                    celda, nombre de hoja, título/mensaje de DV, ítem de lista de DV, cabecera/pie, docProps),
                    su TRATAMIENTO y su GRUPO de traducción (G0-comun, G01…, GL-mercado, GM-mapas).
  · censo_es.json   por libro y hoja, lo que contarán los gates de la F2-EN: fórmulas (con huella sha1),
                    referencias a hojas, literales en fórmulas y en CF, DV, CF, gráficos (del ZIP), merges,
                    protección, paneles, áreas y títulos de impresión, papel, cabeceras/pies, formatos (con
                    «€»), fechas y su numFmt, celdas verdes/desbloqueadas, columnas ocultas, docProps.

    python3 extraer_textos.py [--origen DIR] [--salida DIR]

No traduce nada (regla 1bis: los textos EN los escriben subagentes Anthropic, nunca bridge.py).
Térmica: un solo proceso, en serie, sin builds ni navegador.

Tratamientos (por aparición; la cadena toma el más «fuerte»: traducir > mapa-hoja > mapa-clave > localizar > regenerar):
  traducir    texto libre → lo traduce un subagente (grupos G0-comun y G01…).
  mapa-hoja   nombre de pestaña, o celda cuyo valor ES exactamente una pestaña del libro → `mapas.HOJAS`.
  mapa-clave  categoría, unidad, ítem de DV, familia de recepción, zona, mes → tablas de `mapas.py`.
  localizar   DATO DE EJEMPLO (producto, proveedor y sus fichas, notas «(ejemplo)») → sale de la tabla
              maestra de mercado de la F2 (SPEC D14-D15), no de una traducción literal (grupo GL-mercado).
  regenerar   lo reescribe `aplicar_en.py`: línea de versión, subject/description de docProps, pie
              «Página &P de &N», tabla de temperaturas del 04 (columnas B y E, D12) y excepciones de IVA
              del 03 (D11).
"""
import argparse
import datetime
import hashlib
import json
import math
import os
import posixpath
import re
import subprocess
import sys
import zipfile
from collections import OrderedDict, defaultdict

import openpyxl
from openpyxl.styles.numbers import is_date_format
from openpyxl.utils import get_column_letter, range_boundaries

AQUI = os.path.dirname(os.path.abspath(__file__))
REPO = os.path.abspath(os.path.join(AQUI, '..', '..', '..'))
ORIGEN = os.path.join(REPO, 'astro-site', 'public', 'dl', 'kit-inventario')
sys.path.insert(0, AQUI)
import mapas                                                     # noqa: E402

VERDE = 'E8F5E9'
INVARIABLES = {'OK', 'N/A', 'AI Chef Pro'}          # se quedan igual en inglés (fuera de textos_es)
RX_VERSION = re.compile(r'^Versi[óo]n \d+\.\d+ · ')
RX_LITERAL = re.compile(r'"((?:[^"]|"")*)"')
RX_REF_HOJA = re.compile(r"(?:'((?:[^']|'')+)'|([^\W\d][\w.]*))!(?=\$?[A-Za-z]{1,3}\$?\d|\$?[A-Za-z]{1,3}:)")
RX_CITA = re.compile(r"«([^»]{1,40})»|'([^']{1,40})'|\"([^\"]{1,40})\"|“([^”]{1,40})”|‘([^’]{1,40})’")
TAM_MIN, TAM_OBJ, TAM_MAX = 80, 120, 150

# Pistas de la SPEC que se adjuntan a la cadena (para el traductor). Solo orientan; no traducen.
PISTAS = [
    (r'\bIVA\b|base imponible|[Cc]uota', 'D11: impuesto neutro → «Tax %», «Net (pre-tax) amount», «Tax amount»; '
                                         'importes «net of recoverable tax (ex VAT in the UK)»; nada de tipos españoles'),
    (r'€|\bEUR\b|euros?', 'D13: sin símbolo de moneda; los rótulos pierden «(€)»; importes en «your currency»; '
                         '«$» solo en ejemplos de texto US'),
    (r'APPCC', 'Glosario: APPCC → «HACCP / food safety plan»'),
    (r'RGSEAA', 'D16: Nº RGSEAA → «Food license / FDA reg. no.» (approved source, Food Code §3-201.11)'),
    (r'CIF|NIF', 'D16: CIF/NIF → «Tax ID (EIN / VAT no.)»'),
    (r'homolog', 'Glosario: homologado → «approved» (approved supplier / approved source)'),
    (r'°C|Tª|[Tt]emperatura', 'D12: °F primero y °C entre paréntesis en el texto; los números comparables van en °F; '
                              'límites del FDA Food Code 2022 (tabla FAMILIAS de mapas.py)'),
    (r'Reglamento|Reg\. \(CE\)|\bRD\b|\(UE\)|1169/2011|853/2004|852/2004|589/2008|3484/2000',
     'D12/D17: la norma española/UE se sustituye por FDA Food Code 2022 (§ citado en FAMILIAS); UK solo en nota'),
    (r'[Hh]uev', 'D12: en EE. UU. los huevos se REFRIGERAN (≤ 45 °F, 21 CFR 118.4(e)); lo contrario que el ES'),
    (r'[Cc]aducidad|[Cc]onsumo preferente', 'D18: caducidad → «use-by» (se tira); consumo preferente → «best-by» '
                                           '(se revisa); fecha de entrada → «received date»'),
    (r'[Mm]erma', 'Glosario: merma → «waste» (merma de despiece → «butchery trim loss»)'),
    (r'[Cc]ubierto|[Cc]omensal', 'Glosario: covers (UK igual); ticket medio → «average check»'),
    (r'[Ee]xistencias', 'Glosario: existencias iniciales/finales → «beginning/ending inventory» (UK: opening/closing stock)'),
    (r'albar[áa]n', 'Glosario: albarán → «invoice / packing slip»'),
    (r'[Aa]bono', 'Glosario: abono → «credit memo» / «credit»'),
    (r'[Rr]appel', 'Glosario: rappel → «volume rebate»'),
    (r'[Pp]ortes', 'Glosario: portes → «delivery fee»'),
    (r'[Ee]conomato', 'Glosario: economato → «dry storage»'),
    (r'[Cc]ámara', 'Glosario: cámara → «walk-in cooler»; congelador → «walk-in freezer»'),
    (r'[Pp]roveedor', 'Glosario: proveedor → «vendor» (UK: supplier, solo en notas)'),
    (r'A4', 'D19: papel US Letter («Letter landscape»)'),
    (r'Revisar →|[Dd]esproteger|Archivo →', 'Rutas de Excel en inglés: «Review → Unprotect Sheet», '
                                           '«File → Print → Save as PDF»'),
    (r'[Pp]lantilla \d+|BONUS-0\d|fichero 0\d', 'D5: cita de plantilla → número + título EN (mapas.TITULOS)'),
    (r'[«»]', 'D6: en EN las pestañas y los valores se citan entre comillas dobles rectas ("Current Order")'),
    (r'hasta la( fila)? \d+|filas? \d+', 'Cifra derivada (D20): la fila la recalcula aplicar_en.py'),
    (r'Página', 'D19: «Page &P of &N»'),
    (r'Kit (de )?(Control de )?Inventario', 'D2: «Restaurant Inventory Kit Pro»'),
    (r'aichef\.pro/kit-inventario', 'D21: ' + mapas.URL_PRODUCTO),
    (r'Todos los derechos reservados|©', '«© 2026 AI Chef Pro · aichef.pro»'),
    (r'Productos digitales', 'docProps category = «AI Chef Pro · Digital products»'),
    (r'RT-\d+', 'Defecto heredado I1: ID interno de auditoría visible; en EN se elimina'),
    (r'[Pp]ar [Ll]evel|[Pp]ar [Mm]ax', 'Glosario: par level / par max (término US, se mantiene)'),
    (r'[Pp]unto de pedido', 'Glosario: punto de pedido → «reorder point»; stock de seguridad → «safety stock»'),
]


# --------------------------------------------------------------------------
# utilidades
# --------------------------------------------------------------------------
def sha256(path):
    h = hashlib.sha256()
    with open(path, 'rb') as fh:
        for b in iter(lambda: fh.read(65536), b''):
            h.update(b)
    return h.hexdigest()


def corto(fichero):
    """'01', '02'… y 'B08'/'B09' para los BONUS (el split del piloto los fundía en 'BONUS')."""
    m = re.match(r'^(BONUS-)?(\d+)', fichero)
    return ('B' if m.group(1) else '') + m.group(2)


def tiene_letra(s):
    return any(ch.isalpha() for ch in s)


def literales(formula):
    return [m.group(1).replace('""', '"') for m in RX_LITERAL.finditer(formula or '')]


def refs_hoja(formula):
    sin_lit = RX_LITERAL.sub('""', formula or '')
    out = []
    for m in RX_REF_HOJA.finditer(sin_lit):
        out.append((m.group(1) or m.group(2)).replace("''", "'"))
    return out


def celdas_de(sqref):
    """Conjunto de coordenadas (fila, col) de un sqref (varios rangos separados por espacios)."""
    s = set()
    for rng in str(sqref).split():
        c1, r1, c2, r2 = range_boundaries(rng)
        for r in range(r1, r2 + 1):
            for c in range(c1, c2 + 1):
                s.add((r, c))
    return s


def _xml_attr(tag_xml, attr):
    m = re.search(r'\b' + attr + r'="([^"]*)"', tag_xml)
    return m.group(1) if m else None


def _resolver(base_parte, target):
    if target.startswith('/'):
        return target.lstrip('/')
    return posixpath.normpath(posixpath.join(posixpath.dirname(base_parte), target))


def graficos(path):
    """Gráficos del libro leídos del ZIP (openpyxl los borra al cargar): hoja ancla, tipo, título, refs."""
    out = []
    with zipfile.ZipFile(path) as z:
        nombres = set(z.namelist())
        if not any(n.startswith('xl/charts/') for n in nombres):
            return out
        wbx = z.read('xl/workbook.xml').decode('utf-8')
        rels = z.read('xl/_rels/workbook.xml.rels').decode('utf-8')
        rid2t = {_xml_attr(t, 'Id'): _xml_attr(t, 'Target') for t in re.findall(r'<Relationship\b[^>]*>', rels)}
        hoja_de_parte = {}
        for t in re.findall(r'<sheet\b[^>]*>', wbx):
            nombre = _xml_attr(t, 'name')
            rid = _xml_attr(t, 'r:id')
            parte = _resolver('xl/workbook.xml', rid2t[rid])
            hoja_de_parte[parte] = nombre.replace('&amp;', '&')
        ancla = {}
        for parte, hoja in hoja_de_parte.items():
            srel = posixpath.join(posixpath.dirname(parte), '_rels', posixpath.basename(parte) + '.rels')
            if srel not in nombres:
                continue
            for t in re.findall(r'<Relationship\b[^>]*>', z.read(srel).decode('utf-8')):
                if _xml_attr(t, 'Type').endswith('/drawing'):
                    dparte = _resolver(parte, _xml_attr(t, 'Target'))
                    drel = posixpath.join(posixpath.dirname(dparte), '_rels', posixpath.basename(dparte) + '.rels')
                    if drel in nombres:
                        for t2 in re.findall(r'<Relationship\b[^>]*>', z.read(drel).decode('utf-8')):
                            if _xml_attr(t2, 'Type').endswith('/chart'):
                                ancla[_resolver(dparte, _xml_attr(t2, 'Target'))] = hoja
        for parte in sorted(n for n in nombres if re.fullmatch(r'xl/charts/chart\d+\.xml', n)):
            x = z.read(parte).decode('utf-8')
            tipo = re.findall(r'<(?:c:)?(\w+Chart)>', x)
            titulo = None
            mt = re.search(r'<(?:c:)?title>(.*?)</(?:c:)?title>', x, re.S)
            if mt:
                titulo = ''.join(re.findall(r'<a:t>([^<]*)</a:t>', mt.group(1))) or None
            refs = re.findall(r'<(?:c:)?f>([^<]*)</(?:c:)?f>', x)
            out.append(OrderedDict([
                ('parte', parte), ('hoja', ancla.get(parte)), ('tipos', tipo), ('titulo', titulo),
                ('refs', [r.replace('&apos;', "'").replace('&amp;', '&') for r in refs]),
                ('hojas_citadas', sorted({h for r in refs for h in refs_hoja(r.replace('&apos;', "'"))})),
                ('textos_cache', re.findall(r'<(?:c:)?v>([^<]*[A-Za-zÁ-ú][^<]*)</(?:c:)?v>', x)),
            ]))
    return out


# --------------------------------------------------------------------------
# censo y apariciones de un libro
# --------------------------------------------------------------------------
CLAVES = mapas.todas_las_claves()
RX_EJEMPLO = re.compile(r'\(ejemplo\)\s*$')
# Celdas de ficha de proveedor de ejemplo: se localizan enteras (D15)
ZONAS_LOCALIZAR = {
    ('02', 'Directorio Proveedores'): (4, 9, 'BCDEFGHIJKLOS'),
    ('02', 'Condiciones Comerciales'): (4, 9, 'FGJ'),
    ('03', 'Proveedores'): (4, 9, 'ACDFH'),
    ('03', 'Pedido Actual'): (3, 6, 'C'),
    ('03', 'Historial Pedidos'): (5, 7, 'BCDI'),
    ('04', 'Control Recepción'): (4, 7, 'BCDHTU'),
    ('04', 'Registro Incidencias'): (4, 6, 'BCDFGH'),
    ('05', 'Registro Diario Mermas'): (4, 6, 'BIJ'),
    ('05', 'Plan de Acción'): (4, 8, 'CDEF'),
    ('06', 'Control FIFO'): (5, 10, 'BDST'),
    ('06', 'Mapa Almacén'): (4, 14, 'BCDE'),
    ('07', 'Top 20 Productos'): (4, 23, 'BK'),
    ('B09', 'Calculadora'): (4, 33, 'BO'),
}
# Celdas que reescribe aplicar_en.py (no se traducen): D11 (IVA) y D12 (temperaturas)
ZONAS_REGENERAR = {
    ('03', 'Listas'): (2, 20, 'EF', 'D11 excepciones de impuesto: tabla neutra'),
    ('04', 'Verificación Temperaturas'): (4, 18, 'BE', 'D12 tabla FDA Food Code (mapas.FAMILIAS)'),
}
PRODUCTOS_ES = set()          # se rellena con la columna B de las hojas de producto (01 y BONUS-08)


def base_producto(texto):
    t = RX_EJEMPLO.sub('', texto).strip()
    t = re.sub(r', docena$', '', t)
    return t


def en_zona(zonas, clave, cel):
    z = zonas.get(clave)
    if not z:
        return None
    r1, r2, cols = z[0], z[1], z[2]
    if r1 <= cel.row <= r2 and get_column_letter(cel.column) in cols:
        return z
    return None


def clasificar_celda(fichero, ws, cel, texto, hojas_libro):
    """(tratamiento, detalle) de una celda de texto."""
    f = corto(fichero)
    if RX_VERSION.match(texto):
        return 'regenerar', 'D21 línea de versión'
    z = en_zona(ZONAS_REGENERAR, (f, ws.title), cel)
    if z:
        return 'regenerar', z[3]
    if texto in hojas_libro:
        return 'mapa-hoja', 'celda que cita la pestaña'
    if texto in CLAVES:
        return 'mapa-clave', CLAVES[texto][0]
    if base_producto(texto) in PRODUCTOS_ES:
        return 'localizar', 'producto (tabla maestra D14)'
    if en_zona(ZONAS_LOCALIZAR, (f, ws.title), cel):
        return 'localizar', 'dato de ejemplo (D15)'
    if RX_EJEMPLO.search(texto):
        return 'localizar', 'dato de ejemplo «(ejemplo)» (D15)'
    return 'traducir', None


def unidades_dv(ws):
    """¿La hoja lleva la DV de unidades del kit?"""
    for dv in ws.data_validations.dataValidation:
        f1 = dv.formula1 or ''
        if f1.startswith('"') and set(f1.strip('"').split(',')) == set(mapas.UNIDADES):
            return True
    return False


def censar_libro(path, fichero):
    wb = openpyxl.load_workbook(path)
    hojas = list(wb.sheetnames)
    ocurr = []                      # (texto, dict)
    info = OrderedDict()
    info['sha256'] = sha256(path)
    info['nombre_en'] = mapas.FICHEROS[fichero]
    info['titulo_en'] = mapas.TITULOS[mapas.FICHEROS[fichero]]
    info['hojas'] = hojas
    info['hojas_en'] = [mapas.HOJAS.get(h) for h in hojas]
    info['hojas_estado'] = OrderedDict((ws.title, ws.sheet_state) for ws in wb.worksheets)
    info['nombres_definidos'] = sorted(getattr(wb.defined_names, 'keys', lambda: [])())
    pr = wb.properties
    info['docprops'] = OrderedDict((k, getattr(pr, k)) for k in
                                   ('creator', 'title', 'subject', 'keywords', 'description', 'category',
                                    'lastModifiedBy'))
    info['hojas_detalle'] = OrderedDict()
    fx_hoja_total = 0
    for h in hojas:
        ocurr.append((h, {'f': corto(fichero), 'h': h, 't': 'hoja', 'tr': 'mapa-hoja', 'det': 'nombre de pestaña'}))

    for ws in wb.worksheets:
        hd = OrderedDict()
        hd['nombre_en'] = mapas.HOJAS.get(ws.title)
        hd['dimension'] = ws.dimensions
        hd['max_fila'] = ws.max_row
        hd['max_col'] = ws.max_column
        hd['dv_unidades'] = unidades_dv(ws)
        n_txt = n_num = n_fx = 0
        fx_con_hoja = 0
        refs = defaultdict(int)
        lits = defaultdict(list)
        formatos = defaultdict(int)
        fechas = []
        verdes = verdes_desbl = desbl = 0
        eur = ddmm = 0
        celsius = []
        huella = hashlib.sha1()
        for row in ws.iter_rows():
            for cel in row:
                if cel.__class__.__name__ == 'MergedCell':
                    continue
                v = cel.value
                fmt = cel.number_format
                try:
                    fid = cel._style.numFmtId
                except AttributeError:
                    fid = None
                if is_date_format(fmt):
                    fechas.append(OrderedDict([('celda', cel.coordinate), ('numFmtId', fid), ('formato', fmt),
                                               ('vacia', v is None)]))
                if fmt and fmt != 'General':
                    formatos[fmt] += 1
                if '€' in (fmt or ''):
                    eur += 1
                if 'dd/mm' in (fmt or '').lower():
                    ddmm += 1
                fill = cel.fill.fgColor.rgb if cel.fill is not None and cel.fill.fgColor is not None else None
                bloq = cel.protection is None or cel.protection.locked is not False
                if not bloq:
                    desbl += 1
                if isinstance(fill, str) and fill.upper().endswith(VERDE):
                    verdes += 1
                    if not bloq:
                        verdes_desbl += 1
                if v is None:
                    continue
                if isinstance(v, str) and cel.data_type == 'f':
                    n_fx += 1
                    huella.update('{}={}\n'.format(cel.coordinate, v).encode('utf-8'))
                    for lit in literales(v):
                        lits[lit].append(cel.coordinate)
                    rh = refs_hoja(v)
                    if rh:
                        fx_con_hoja += 1
                    for nom in rh:
                        refs[nom] += 1
                    continue
                if isinstance(v, str):
                    n_txt += 1
                    if '°C' in v or 'Tª' in v:
                        celsius.append(cel.coordinate)
                    tr, det = clasificar_celda(fichero, ws, cel, v, hojas)
                    d = {'f': corto(fichero), 'h': ws.title, 'c': cel.coordinate, 't': 'celda', 'tr': tr}
                    if det:
                        d['det'] = det
                    ocurr.append((v, d))
                else:
                    n_num += 1
        fx_hoja_total += fx_con_hoja
        hd['celdas_texto'] = n_txt
        hd['celdas_numero'] = n_num
        hd['celdas_formula'] = n_fx
        hd['formulas_con_hoja'] = fx_con_hoja
        hd['refs_hoja'] = OrderedDict(sorted(refs.items()))
        hd['literales'] = OrderedDict((k, lits[k]) for k in sorted(lits))
        hd['literales_recuento'] = OrderedDict((k, len(lits[k])) for k in sorted(lits))
        hd['huella_formulas_sha1'] = huella.hexdigest() if n_fx else None
        hd['fechas'] = fechas
        hd['formatos'] = OrderedDict(sorted(formatos.items(), key=lambda kv: -kv[1]))
        hd['celdas_formato_eur'] = eur          # todas las celdas (también vacías) con «€» en el formato
        hd['celdas_formato_ddmm'] = ddmm        # ídem con «dd/mm»
        hd['textos_celsius'] = celsius          # celdas de texto con «°C» o «Tª» (D12)
        hd['verdes'] = verdes
        hd['verdes_desbloqueadas'] = verdes_desbl
        hd['desbloqueadas'] = desbl
        hd['columnas_ocultas'] = sorted(k for k, d in ws.column_dimensions.items() if d.hidden)
        hd['filas_ocultas'] = sorted(k for k, d in ws.row_dimensions.items() if d.hidden)
        hd['merges'] = sorted(str(m) for m in ws.merged_cells.ranges)
        hd['paneles'] = ws.freeze_panes
        p = ws.protection
        hd['proteccion'] = OrderedDict([('hoja', bool(p.sheet)), ('password', bool(p.password)),
                                        ('sort', p.sort), ('autoFilter', p.autoFilter),
                                        ('formatColumns', p.formatColumns), ('formatRows', p.formatRows),
                                        ('insertRows', p.insertRows), ('deleteRows', p.deleteRows),
                                        ('selectLockedCells', p.selectLockedCells),
                                        ('selectUnlockedCells', p.selectUnlockedCells)])
        hd['autofiltro'] = ws.auto_filter.ref
        hd['area_impresion'] = ws.print_area
        hd['titulos_impresion'] = OrderedDict([('filas', ws.print_title_rows), ('columnas', ws.print_title_cols)])
        ps = ws.page_setup
        hd['papel'] = OrderedDict([('paperSize', ps.paperSize), ('orientation', ps.orientation),
                                   ('fitToWidth', ps.fitToWidth), ('fitToHeight', ps.fitToHeight),
                                   ('fitToPage', bool(ws.sheet_properties.pageSetUpPr and
                                                      ws.sheet_properties.pageSetUpPr.fitToPage))])
        hf = OrderedDict()
        for nombre in ('oddHeader', 'oddFooter', 'evenHeader', 'evenFooter', 'firstHeader', 'firstFooter'):
            obj = getattr(ws, nombre)
            for parte in ('left', 'center', 'right'):
                txt = getattr(obj, parte).text
                if txt:
                    hf['{}.{}'.format(nombre, parte)] = txt
                    ocurr.append((txt, {'f': corto(fichero), 'h': ws.title, 't': 'cabecera-pie',
                                        'c': '{}.{}'.format(nombre, parte), 'tr': 'regenerar',
                                        'det': 'D19 pie «AI Chef Pro · aichef.pro · Page &P of &N»'}))
        hd['cabeceras_pies'] = hf
        # DV
        dvs = []
        for dv in ws.data_validations.dataValidation:
            dvs.append(OrderedDict([
                ('sqref', str(dv.sqref)), ('type', dv.type), ('operator', dv.operator),
                ('formula1', dv.formula1), ('formula2', dv.formula2), ('allow_blank', dv.allow_blank),
                ('showDropDown', dv.showDropDown), ('showErrorMessage', dv.showErrorMessage),
                ('showInputMessage', dv.showInputMessage), ('errorStyle', dv.errorStyle),
                ('errorTitle', dv.errorTitle), ('error', dv.error),
                ('promptTitle', dv.promptTitle), ('prompt', dv.prompt),
                ('n_celdas', len(celdas_de(dv.sqref))),
                ('cita_hojas', sorted(set(refs_hoja(dv.formula1 or '')))),
            ]))
            base = {'f': corto(fichero), 'h': ws.title, 'c': str(dv.sqref)}
            for campo, t in (('errorTitle', 'dv-error-titulo'), ('error', 'dv-error'),
                             ('promptTitle', 'dv-prompt-titulo'), ('prompt', 'dv-prompt')):
                txt = getattr(dv, campo)
                if txt:
                    ocurr.append((txt, dict(base, t=t, tr='traducir')))
            f1 = dv.formula1 or ''
            if f1.startswith('"'):
                for it in f1.strip('"').split(','):
                    if it in CLAVES:
                        d = dict(base, t='dv-item', tr='mapa-clave', det=CLAVES[it][0])
                    elif re.fullmatch(r'\d+', it):
                        d = dict(base, t='dv-item', tr='regenerar', det='D11 lista de tipos de IVA → DV decimal')
                    else:
                        d = dict(base, t='dv-item', tr='traducir')
                    ocurr.append((it, d))
        hd['dv'] = dvs
        # CF
        cfs = []
        lits_cf = defaultdict(int)
        for rng, reglas in ws.conditional_formatting._cf_rules.items():
            for rg in reglas:
                fs = list(rg.formula or [])
                cfs.append(OrderedDict([
                    ('sqref', str(rng.sqref)), ('type', rg.type), ('operator', rg.operator),
                    ('text', getattr(rg, 'text', None)),
                    ('formula', fs), ('priority', rg.priority), ('stopIfTrue', rg.stopIfTrue),
                    ('cita_hojas', sorted({h for f in fs for h in refs_hoja(f)})),
                ]))
                for f in fs:
                    for lit in literales(f):
                        lits_cf[lit] += 1
        hd['cf'] = cfs
        hd['literales_cf'] = OrderedDict(sorted(lits_cf.items()))
        info['hojas_detalle'][ws.title] = hd

    # gráficos e imágenes (del ZIP)
    info['graficos'] = graficos(path)
    with zipfile.ZipFile(path) as z:
        info['imagenes'] = sorted(n for n in z.namelist() if n.startswith('xl/media/'))
    for g in info['graficos']:
        if g['titulo']:
            ocurr.append((g['titulo'], {'f': corto(fichero), 'h': g['hoja'], 't': 'grafico-titulo',
                                        'c': g['parte'], 'tr': 'traducir'}))
    # docProps
    for campo, v in info['docprops'].items():
        if not isinstance(v, str) or not v:
            continue
        if campo == 'subject':
            tr, det = 'regenerar', 'D21 subject «%s · v%s»' % (mapas.PRODUCTO, mapas.VERSION_ES)
        elif campo == 'description':
            tr, det = 'regenerar', 'D21 URL ' + mapas.URL_PRODUCTO
        elif campo == 'title':
            tr, det = 'regenerar', 'D5 título EN (mapas.TITULOS) · ' + mapas.PRODUCTO
        elif campo == 'category':
            tr, det = 'regenerar', 'AI Chef Pro · Digital products'
        else:
            tr, det = 'traducir', None
        d = {'f': corto(fichero), 't': 'docprops', 'c': campo, 'tr': tr}
        if det:
            d['det'] = det
        ocurr.append((v, d))

    # totales del libro
    hdv = info['hojas_detalle'].values()
    lit_tot = defaultdict(int)
    for x in hdv:
        for k, v in x['literales_recuento'].items():
            lit_tot[k] += v
    litcf_tot = defaultdict(int)
    for x in hdv:
        for k, v in x['literales_cf'].items():
            litcf_tot[k] += v
    info['totales'] = OrderedDict([
        ('hojas', len(hojas)),
        ('celdas_texto', sum(x['celdas_texto'] for x in hdv)),
        ('celdas_formula', sum(x['celdas_formula'] for x in hdv)),
        ('formulas_con_hoja', fx_hoja_total),
        ('dv', sum(len(x['dv']) for x in hdv)),
        ('dv_lista_literal', sum(1 for x in hdv for d in x['dv'] if (d['formula1'] or '').startswith('"'))),
        ('dv_cita_hoja', sum(1 for x in hdv for d in x['dv'] if d['cita_hojas'])),
        ('dv_celdas', sum(d['n_celdas'] for x in hdv for d in x['dv'])),
        ('cf', sum(len(x['cf']) for x in hdv)),
        ('cf_cita_hoja', sum(1 for x in hdv for d in x['cf'] if d['cita_hojas'])),
        ('graficos', len(info['graficos'])),
        ('imagenes', len(info['imagenes'])),
        ('merges', sum(len(x['merges']) for x in hdv)),
        ('hojas_protegidas', sum(1 for x in hdv if x['proteccion']['hoja'])),
        ('areas_impresion', sum(1 for x in hdv if x['area_impresion'])),
        ('titulos_impresion', sum(1 for x in hdv if x['titulos_impresion']['filas'])),
        ('autofiltros', sum(1 for x in hdv if x['autofiltro'])),
        ('fechas', sum(len(x['fechas']) for x in hdv)),
        ('fechas_numFmtId_14', sum(1 for x in hdv for f in x['fechas'] if f['numFmtId'] == 14)),
        ('celdas_formato_eur', sum(x['celdas_formato_eur'] for x in hdv)),
        ('celdas_formato_ddmm', sum(x['celdas_formato_ddmm'] for x in hdv)),
        ('textos_celsius', sum(len(x['textos_celsius']) for x in hdv)),
        ('verdes', sum(x['verdes'] for x in hdv)),
        ('verdes_desbloqueadas', sum(x['verdes_desbloqueadas'] for x in hdv)),
        ('desbloqueadas', sum(x['desbloqueadas'] for x in hdv)),
        ('columnas_ocultas', sum(len(x['columnas_ocultas']) for x in hdv)),
        ('hojas_dv_unidades', sum(1 for x in hdv if x['dv_unidades'])),
        ('papel', sorted({x['papel']['paperSize'] for x in hdv if x['papel']['paperSize'] is not None})),
        ('literales', OrderedDict(sorted(lit_tot.items()))),
        ('literales_cf', OrderedDict(sorted(litcf_tot.items()))),
    ])
    return info, ocurr


# --------------------------------------------------------------------------
# cadenas, grupos y pistas
# --------------------------------------------------------------------------
PRIORIDAD = {'traducir': 0, 'mapa-hoja': 1, 'mapa-clave': 2, 'localizar': 3, 'regenerar': 4}


def en_de_mapa(texto, dets):
    if texto in mapas.HOJAS and 'mapa-hoja' in dets:
        return mapas.HOJAS[texto]
    if texto in CLAVES:
        return CLAVES[texto][1]
    if texto in mapas.HOJAS:
        return mapas.HOJAS[texto]
    return None


def pistas(texto, fila_rango):
    out = []
    for rx, nota in PISTAS:
        if re.search(rx, texto):
            out.append(nota.format(fila=fila_rango))
    return out


def citas_hoja(texto, todas_hojas):
    out = []
    for m in RX_CITA.finditer(texto):
        span = next(g for g in m.groups() if g is not None)
        if span in todas_hojas and span not in [c['es'] for c in out]:
            out.append(OrderedDict([('es', span), ('en', mapas.HOJAS.get(span))]))
    return out


def categorias_citadas(texto):
    out = []
    for es, en in mapas.CATEGORIAS.items():
        if re.search(r'(?<!\w)' + re.escape(es) + r'(?!\w)', texto):
            out.append(OrderedDict([('es', es), ('en', en)]))
    return out


def palabras(s):
    return len(re.findall(r'\w+', s))


def agrupar(orden_ficheros, por_fichero):
    """Bloques de ~80-150 cadenas por fichero (juntando ficheros pequeños)."""
    grupos = []
    cur, cur_f = [], []

    def cerrar():
        nonlocal cur, cur_f
        if cur:
            grupos.append((list(cur_f), list(cur)))
        cur, cur_f = [], []

    for f in orden_ficheros:
        ids = por_fichero.get(f, [])
        if not ids:
            continue
        n = len(ids)
        if n > TAM_MAX:
            cerrar()
            k = int(math.ceil(n / float(TAM_OBJ)))
            tam = int(math.ceil(n / float(k)))
            for i in range(k):
                grupos.append(([f + ' ({}/{})'.format(i + 1, k)], ids[i * tam:(i + 1) * tam]))
            continue
        if cur and len(cur) + n > TAM_MAX:
            cerrar()
        cur += ids
        cur_f.append(f)
        if len(cur) >= TAM_MIN:
            cerrar()
    if cur:
        if grupos and len(cur) < TAM_MIN // 2 and len(grupos[-1][1]) + len(cur) <= TAM_MAX + 30 \
                and '(' not in grupos[-1][0][-1]:
            grupos[-1] = (grupos[-1][0] + cur_f, grupos[-1][1] + cur)
        else:
            grupos.append((cur_f, cur))
    return grupos


def main():
    ap = argparse.ArgumentParser(description=__doc__.split('\n')[1])
    ap.add_argument('--origen', default=ORIGEN)
    ap.add_argument('--salida', default=AQUI)
    args = ap.parse_args()

    libros = [f for f in mapas.LIBROS_ES]
    # productos de ejemplo = columna B de las hojas de producto del 01 y del BONUS-08 (D14)
    for f, hojas_p in (('01-inventario-stock-diario.xlsx', ('Cocina', 'Barra', 'Almacén')),
                       ('BONUS-08-inventario-rapido-mensual.xlsx', ('Conteo Rápido',))):
        pth = os.path.join(args.origen, f)
        if os.path.exists(pth):
            wbp = openpyxl.load_workbook(pth)
            for h in hojas_p:
                wsp = wbp[h]
                for r in range(5, wsp.max_row + 1):
                    v = wsp.cell(row=r, column=2).value
                    if isinstance(v, str) and v.strip():
                        PRODUCTOS_ES.add(v.strip())
    PRODUCTOS_ES.update({'Merluza fresca'})      # único producto de ejemplo de 04/05/06 fuera de la lista
    faltan = [f for f in libros if not os.path.exists(os.path.join(args.origen, f))]
    if faltan:
        sys.exit('faltan ficheros en {}: {}'.format(args.origen, faltan))
    sobran = sorted(f for f in os.listdir(args.origen) if f.endswith('.xlsx') and f not in libros)
    if sobran:
        sys.exit('xlsx no previstos en {}: {}'.format(args.origen, sobran))

    try:
        commit = subprocess.check_output(['git', '-C', REPO, 'rev-parse', '--short', 'HEAD'],
                                         stderr=subprocess.DEVNULL).decode().strip()
    except Exception:                                             # noqa: BLE001
        commit = None

    censo = OrderedDict()
    censo['meta'] = OrderedDict([
        ('generado', datetime.datetime.utcnow().replace(microsecond=0).isoformat() + 'Z'),
        ('script', 'scripts/productos-digitales/restaurant-inventory-kit/extraer_textos.py'),
        ('origen', os.path.relpath(args.origen, REPO)),
        ('commit', commit),
        ('nota', 'Censo del Kit de Inventario v2.0 PUBLICADO (SPEC §4). Los gates de la F2-EN comparan '
                 'contra este fichero, no contra recuentos fijos. Fechas = celdas con formato de fecha '
                 '(en EN pasan a numFmtId 14 = fecha corta del sistema, D19).'),
    ])
    censo['libros'] = OrderedDict()
    todas = []                       # (texto, ocurrencia)
    for f in libros:
        info, oc = censar_libro(os.path.join(args.origen, f), f)
        censo['libros'][f] = info
        todas.extend(oc)
        print('  {:<40} {:>2} hojas · {:>4} textos · {:>4} fórmulas'.format(
            f, info['totales']['hojas'], info['totales']['celdas_texto'], info['totales']['celdas_formula']))

    # totales del kit
    T = OrderedDict()
    for k in ('hojas', 'celdas_texto', 'celdas_formula', 'formulas_con_hoja', 'dv', 'dv_lista_literal',
              'dv_cita_hoja', 'dv_celdas', 'cf', 'cf_cita_hoja', 'graficos', 'imagenes', 'merges',
              'hojas_protegidas', 'areas_impresion', 'titulos_impresion', 'autofiltros', 'fechas',
              'fechas_numFmtId_14', 'celdas_formato_eur', 'celdas_formato_ddmm', 'textos_celsius', 'verdes',
              'verdes_desbloqueadas', 'desbloqueadas', 'columnas_ocultas', 'hojas_dv_unidades'):
        T[k] = sum(censo['libros'][f]['totales'][k] for f in libros)
    lit = defaultdict(int)
    litcf = defaultdict(int)
    for f in libros:
        for k, v in censo['libros'][f]['totales']['literales'].items():
            lit[k] += v
        for k, v in censo['libros'][f]['totales']['literales_cf'].items():
            litcf[k] += v
    T['literales'] = OrderedDict(sorted(lit.items(), key=lambda kv: -kv[1]))
    T['literales_cf'] = OrderedDict(sorted(litcf.items(), key=lambda kv: -kv[1]))
    T['papel'] = sorted({p for f in libros for p in censo['libros'][f]['totales']['papel']})
    censo['totales'] = T

    # ---- cadenas únicas
    todas_hojas = set(h for f in libros for h in censo['libros'][f]['hojas'])
    fila_rango = None
    excl_sin_letra = OrderedDict()
    excl_inv = OrderedDict()
    cad = OrderedDict()
    for texto, oc in todas:
        if not tiene_letra(texto):
            excl_sin_letra[texto] = excl_sin_letra.get(texto, 0) + 1
            continue
        if texto in INVARIABLES:
            excl_inv[texto] = excl_inv.get(texto, 0) + 1
            continue
        cad.setdefault(texto, []).append(oc)

    orden_f = [corto(f) for f in libros]
    cadenas = []
    for i, (texto, ocs) in enumerate(cad.items(), 1):
        fich = []
        for o in ocs:
            if o['f'] not in fich:
                fich.append(o['f'])
        trs = sorted({o['tr'] for o in ocs}, key=lambda t: PRIORIDAD[t])
        tr = trs[0]
        dets = {o.get('det') for o in ocs} | {o['tr'] for o in ocs}
        e = OrderedDict()
        e['id'] = 'c{:04d}'.format(i)
        e['es'] = texto
        e['grupo'] = None
        e['tratamiento'] = tr
        if len(trs) > 1:
            e['tratamientos'] = trs
        enm = en_de_mapa(texto, dets) if (tr != 'traducir' or len(trs) > 1) else None
        if enm is not None:
            e['en_mapa'] = enm
        e['n_ficheros'] = len(fich)
        e['n_apariciones'] = len(ocs)
        e['palabras'] = palabras(texto)
        e['con_cifra'] = bool(re.search(r'\d', texto))
        ch = citas_hoja(texto, todas_hojas)
        if ch:
            e['cita_hojas'] = ch
        if tr == 'traducir':
            cc = categorias_citadas(texto)
            if cc and texto not in mapas.CATEGORIAS:
                e['categorias_citadas'] = cc
            pi = pistas(texto, fila_rango)
            if pi:
                e['pistas'] = pi
        e['sha1'] = hashlib.sha1(texto.encode('utf-8')).hexdigest()[:12]
        e['donde'] = ocs
        e['_ficheros'] = fich
        cadenas.append(e)

    # ---- grupos
    por_fichero = OrderedDict((f, []) for f in orden_f)
    for e in cadenas:
        if e['tratamiento'] == 'localizar':
            e['grupo'] = 'GL-mercado'
            continue
        if e['tratamiento'] != 'traducir':
            e['grupo'] = 'GM-mapas'
            continue
        n = e['n_ficheros']
        instr = any(o.get('h') == 'Instrucciones' for o in e['donde'])
        comun_meta = any(o['t'] == 'docprops' and o['c'] == 'keywords' for o in e['donde'])
        if n >= 3 or (n == 2 and instr) or comun_meta:
            e['grupo'] = 'G0-comun'
        else:
            por_fichero[e['_ficheros'][0]].append(e['id'])
    bloques = agrupar(orden_f, por_fichero)
    idx = {e['id']: e for e in cadenas}
    nombres_f = {corto(f): f for f in libros}
    grupos = []
    g0 = [e for e in cadenas if e['grupo'] == 'G0-comun']
    grupos.append(OrderedDict([
        ('id', 'G0-comun'),
        ('descripcion', 'Común: cadenas en ≥ 3 libros (mensajes de DV, cabeceras repetidas), Instrucciones '
                        'comunes a 2 libros y keywords. Se traduce PRIMERO: fija el glosario.'),
        ('ficheros', sorted({f for e in g0 for f in e['_ficheros']}, key=orden_f.index)),
        ('n_cadenas', len(g0)), ('n_palabras', sum(e['palabras'] for e in g0)),
    ]))
    for k, (fs, ids) in enumerate(bloques, 1):
        gid = 'G{:02d}'.format(k)
        for i in ids:
            idx[i]['grupo'] = gid
        etiquetas = []
        for f in fs:
            m = re.match(r'^(\S+)(.*)$', f)
            etiquetas.append(nombres_f[m.group(1)] + m.group(2))
        grupos.append(OrderedDict([
            ('id', gid),
            ('descripcion', 'Por fichero: ' + ' + '.join(etiquetas)),
            ('ficheros', [re.match(r'^(\S+)', f).group(1) for f in fs]),
            ('n_cadenas', len(ids)), ('n_palabras', sum(idx[i]['palabras'] for i in ids)),
        ]))
    gl = [e for e in cadenas if e['grupo'] == 'GL-mercado']
    grupos.append(OrderedDict([
        ('id', 'GL-mercado'),
        ('descripcion', 'Datos de ejemplo (productos, fichas de proveedor, notas «(ejemplo)», zonas): los '
                        'escribe el subagente de mercado con la tabla maestra (SPEC D14-D15), no se traducen '
                        'palabra a palabra.'),
        ('ficheros', sorted({f for e in gl for f in e['_ficheros']}, key=orden_f.index)),
        ('n_cadenas', len(gl)), ('n_palabras', sum(e['palabras'] for e in gl)),
    ]))
    gm = [e for e in cadenas if e['grupo'] == 'GM-mapas']
    por_tr = defaultdict(int)
    for e in gm:
        por_tr[e['tratamiento']] += 1
    grupos.append(OrderedDict([
        ('id', 'GM-mapas'),
        ('descripcion', 'NO se traduce: pestañas y claves-dato (mapas.py), versión/subject/URL/título '
                        '(D21, D5), pies (D19), tabla de temperaturas del 04 (D12) e impuestos del 03 (D11). ' +
                        ', '.join('{} {}'.format(v, k) for k, v in sorted(por_tr.items()))),
        ('ficheros', sorted({f for e in gm for f in e['_ficheros']}, key=orden_f.index)),
        ('n_cadenas', len(gm)), ('n_palabras', sum(e['palabras'] for e in gm)),
    ]))
    # orden final: por grupo y, dentro, por primera aparición
    orden_g = {g['id']: k for k, g in enumerate(grupos)}
    cadenas.sort(key=lambda e: (orden_g[e['grupo']], int(e['id'][1:])))
    for e in cadenas:
        e.pop('_ficheros')
        for o in e['donde']:
            o.pop('rejilla', None)

    assert all(e['grupo'] for e in cadenas)
    n_trad = sum(1 for e in cadenas if e['tratamiento'] == 'traducir')
    textos = OrderedDict()
    textos['meta'] = OrderedDict([
        ('generado', censo['meta']['generado']),
        ('script', 'scripts/productos-digitales/restaurant-inventory-kit/extraer_textos.py'),
        ('origen', censo['meta']['origen']),
        ('commit', commit),
        ('libros', OrderedDict((corto(f), OrderedDict([('es', f), ('en', mapas.FICHEROS[f]),
                                                      ('sha256', censo['libros'][f]['sha256'])]))
                               for f in libros)),
        ('n_cadenas', len(cadenas)),
        ('n_traducir', n_trad),
        ('n_palabras_traducir', sum(e['palabras'] for e in cadenas if e['tratamiento'] == 'traducir')),
        ('n_apariciones', sum(e['n_apariciones'] for e in cadenas)),
        ('grupos', grupos),
        ('excluidas', OrderedDict([
            ('sin_letras', excl_sin_letra),
            ('invariables', excl_inv),
            ('regla', 'Fuera: números, fórmulas, cadenas sin ninguna letra (→, ✓, —, ?…) y las invariables '
                      '(OK, N/A, AI Chef Pro). Los literales de fórmula están en censo_es.json, no aquí.'),
        ])),
        ('campos', OrderedDict([
            ('donde', 'f = libro (clave de meta.libros), h = hoja, c = celda / sqref de la DV / campo de '
                      'cabecera-pie / parte del gráfico / campo de docProps; t = tipo de lugar (celda, hoja, '
                      'dv-error, dv-error-titulo, dv-prompt, dv-prompt-titulo, dv-item, grafico-titulo, '
                      'cabecera-pie, docprops); tr = tratamiento de esa aparición; det = detalle'),
            ('tratamiento', 'traducir | mapa-hoja | mapa-clave | localizar | regenerar (ver docstring del script)'),
            ('en_mapa', 'valor EN obligatorio que dicta mapas.py (pestaña o clave-dato); si la cadena '
                        'además se traduce, la traducción TIENE que usar ese valor'),
            ('cita_hojas', 'pestañas ES citadas entre comillas en el texto, con su nombre EN (D9: en EN, '
                           'entre comillas dobles rectas)'),
            ('categorias_citadas', 'categorías del kit citadas en el texto, con su nombre EN (D8)'),
            ('con_cifra', 'el texto lleva cifras: si la versión EN cambia la cifra (precio, IVA, °C, fila, '
                          'unidad…), se reescribe con la cifra EN (D20), no se traduce a ciegas'),
            ('pistas', 'recordatorios de la SPEC que aplican a esa cadena (no traducen)'),
        ])),
    ])
    textos['cadenas'] = cadenas

    os.makedirs(args.salida, exist_ok=True)
    p_t = os.path.join(args.salida, 'textos_es.json')
    p_c = os.path.join(args.salida, 'censo_es.json')
    with open(p_t, 'w', encoding='utf-8') as fh:
        json.dump(textos, fh, ensure_ascii=False, indent=1)
        fh.write('\n')
    with open(p_c, 'w', encoding='utf-8') as fh:
        json.dump(censo, fh, ensure_ascii=False, indent=1)
        fh.write('\n')

    print('\ntextos_es.json: {} cadenas únicas ({} a traducir, {} palabras) · {} apariciones'.format(
        len(cadenas), n_trad, textos['meta']['n_palabras_traducir'], textos['meta']['n_apariciones']))
    for g in grupos:
        print('  {:<9} {:>4} cadenas {:>5} palabras  {}'.format(g['id'], g['n_cadenas'], g['n_palabras'],
                                                             g['descripcion'][:90]))
    print('  excluidas: sin letras {} ({} apariciones) · invariables {}'.format(
        len(excl_sin_letra), sum(excl_sin_letra.values()), dict(excl_inv)))
    print('censo_es.json: {}'.format(json.dumps(
        {k: v for k, v in T.items() if not isinstance(v, dict)}, ensure_ascii=False)))
    print('  literales: {}'.format(dict(T['literales'])))
    print('  literales CF: {}'.format(dict(T['literales_cf'])))


if __name__ == '__main__':
    main()
