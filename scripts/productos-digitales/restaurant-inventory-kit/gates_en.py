#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
gates_en.py — Gates F2 del Restaurant Inventory Kit Pro (EN), SPEC §8 (G1-G8), sobre la salida de aplicar_en.py.
Copia adaptada de `recipe-costing-kit/gates_en.py`. Solo LEE: los 9 xlsx EN, los 9 xlsx ES publicados
(dl/kit-inventario, solo lectura), censo_es.json, textos_es.json y mercado_en.json. Ningún recuento está escrito a
mano: todo sale de censo_es.json, de los xlsx ES o de mercado_en.json / mapas.py.

    python3 gates_en.py                                  # contra astro-site/public/dl/restaurant-inventory-templates/
    python3 gates_en.py --dir <scratchpad>/rik-dryrun    # contra un dry-run
    python3 gates_en.py --solo G1,G2 [--mes September] [--json informe.json]
    python3 gates_en.py --autotest --dir …               # inyecta defectos en una COPIA y exige que se detecten

Gates (SPEC §8):
  G1 paridad estructural EN↔ES (hojas por el mapa; fórmula a fórmula con los mapas de hoja y literales, +19 de D25;
     referencias a hoja; DV —tipo igual salvo D11— y CF; merges; protección sin contraseña; paneles; áreas y títulos
     de impresión; autofiltro; columnas ocultas; verdes = desbloqueadas (+3 de D11); 0 gráficos, 0 imágenes);
  G2 restos: cero español, cero «€»/«EUR», cero caracteres fuera de la lista blanca, cero IDs internos, en valores,
     formatos, pestañas, DV, CF, literales, cabeceras/pies y docProps;
  G3 formato: Letter en las 38 hojas; fechas del censo (+D25) con numFmtId 14; cero dd/mm; pie D19; versión D21;
     docProps D21;
  G4 claves: 10 categorías en las DV literales, en las etiquetas de agregación y en los ejemplos; lista de unidades
     D9 donde el ES tenía la suya; todo valor de ejemplo dentro de su lista de DV; familias y zonas = mapas.py;
  G5 CF y literales: cada token de CF casa con un solo estado de su rango; "Y", "Best-by", "Full year" y Jan…Dec
     coinciden con sus DV; comodines de COUNTIF con su estado;
  G6 cálculo: caché en toda fórmula (<v> en el XML), cero errores; historia del 04 ACCEPT / REJECT (too warm) /
     ACCEPT / N/A; food cost del 07 dentro del objetivo; el 06 enseña los 5 estados hoy;
  G7 identidades D14: validar_mercado.py en verde; cada celda de mercado_en.json «libros» escrita tal cual; BONUS-09
     punto de pedido = par del 01 y máx. = par max; importes = cantidad × precio (valores de la caché); cifras D20;
  G8 censo-entregables.py --only <carpeta> --fail (Letter admitido) en 0 defectos.
Sale con código 1 si algún gate falla.
"""
import argparse
import copy
import datetime
import html
import json
import logging
import os
import re
import shutil
import subprocess
import sys
import tempfile
import zipfile
from collections import OrderedDict, Counter, defaultdict

AQUI = os.path.dirname(os.path.abspath(__file__))
SCRIPTS = os.path.dirname(AQUI)
REPO = os.path.dirname(os.path.dirname(SCRIPTS))
sys.path.insert(0, AQUI)

import openpyxl                                          # noqa: E402

import mapas                                             # noqa: E402
import aplicar_en as A                                   # noqa: E402

logging.disable(logging.CRITICAL)

ORIGEN = A.ORIGEN
CENSO_SCRIPT = os.path.join(SCRIPTS, 'censo-entregables.py')
VALIDAR_MERCADO = os.path.join(AQUI, 'validar_mercado.py')
TOL = 0.006
ERRORES = ('#REF!', '#VALUE!', '#NAME?', '#DIV/0!', '#N/A', '#NUM!', '#NULL!')

# Lista blanca de caracteres no ASCII (SPEC G2: ✓ ✗ ⚠ ⛔ 🔴 🟡 🟢 + tipografía del kit). Todo lo demás falla.
SIMBOLOS_OK = set('✓✗⚠⛔—·°§→©…×≤≥−½') | {'\U0001F534', '\U0001F7E1', '\U0001F7E2', '️'}
RX_ID_INTERNO = re.compile(r'\b(?:kitinv-[a-z0-9]+|grupo_[a-z0-9]+|RT-\d+)\b|\[(?:fuente|estimado|derivado)\]|'
                           r'\bSPEC\b|\((?:D|I)\d{1,2}\)')
RX_NORMA_ES = re.compile(r'\b(?:Reg\.\s*\((?:CE|UE)\)|RD \d+/\d{4}|RGSEAA|APPCC|CIF|NIF)\b')
SIGNOS_ES = set('¿¡«»€ºª')
RX_TILDE = re.compile(r'[áéíóúüñÁÉÍÓÚÜÑ]')
LISTA_BLANCA_TILDE = {'café', 'purée', 'sauté', 'sautéed', 'jalapeño', 'jalapeños', 'crème', 'brûlée'}
FUNCIONALES_ES = {
    'de', 'del', 'las', 'el', 'los', 'y', 'que', 'con', 'para', 'por', 'una', 'unos', 'unas', 'es', 'al', 'se',
    'su', 'sus', 'como', 'pero', 'muy', 'cuando', 'donde', 'porque', 'este', 'esta', 'esto', 'estos', 'desde',
    'hasta', 'sobre', 'entre', 'tu', 'tus', 'te', 'mis', 'hoja', 'pestaña', 'celda', 'celdas', 'sin', 'ejemplo',
}
OFICIO_ES = {
    'iva', 'merma', 'mermas', 'albaran', 'albaranes', 'proveedor', 'proveedores', 'cocina', 'ventas', 'compras',
    'compra', 'coste', 'costes', 'plantilla', 'plantillas', 'solomillo', 'merluza', 'inventario', 'factura',
    'desproteger', 'revisar', 'cantidad', 'precio', 'precios', 'unidad', 'unidades', 'docena', 'botella',
    'objetivo', 'resumen', 'instrucciones', 'categoria', 'bebida', 'bebidas', 'mantequilla', 'harina', 'huevos',
    'nata', 'aceite', 'pescado', 'carne', 'verdura', 'verduras', 'alerta', 'revisa', 'importe', 'caducidad',
    'caducado', 'pedido', 'pedidos', 'almacen', 'camara', 'congelador', 'barra', 'recepcion', 'familia', 'lote',
    'homologado', 'abono', 'rappel', 'portes', 'cubiertos', 'economato', 'huevera', 'limpieza', 'secos', 'granos',
    'lacteos', 'carnicos', 'pescados', 'congelados', 'otros', 'saco', 'caja', 'rollo', 'paquete', 'barril', 'ud',
    'pendiente', 'borrador', 'enviado', 'recibido', 'facturado', 'abierta', 'cerrada', 'reclamada', 'alta', 'media',
    'baja', 'ene', 'abr', 'ago', 'dic', 'enero', 'agosto', 'año', 'mes', 'semana', 'fecha', 'nota', 'notas',
    'zona', 'estado', 'total', 'valor', 'bajo', 'pedir', 'conforme', 'rechazar', 'urgente', 'proximo', 'vigente',
}
OFICIO_ES -= {'total', 'media', 'alta', 'nota', 'notas', 'mes', 'caja'}   # también inglés o marca
RX_URL = re.compile(r'https?://\S+|\b[\w-]+(?:\.[\w-]+)*\.(?:es|com|pro|gov|org|uk|co|example)\b[/\w.#-]*', re.I)
RX_PALABRA = re.compile(r"[A-Za-zÀ-ÖØ-öø-ÿ]+")
RX_STR = A.RX_STR


def no_permitidos(t):
    return sorted({ch for ch in t if ord(ch) > 127 and ch not in SIMBOLOS_OK
                   and not ('À' <= ch <= 'ÿ' and ch not in '×÷')})


def restos_espanol(t):
    out = []
    limpio = RX_URL.sub(' ', t)
    for ch in limpio:
        if ch in SIGNOS_ES:
            out.append(('signo «%s»' % ch, t[:80]))
    for m in RX_PALABRA.finditer(limpio):
        w = m.group(0)
        wl = w.lower()
        ctx = limpio[max(0, m.start() - 30):m.end() + 30].replace('\n', ' ')
        sin = (wl.replace('á', 'a').replace('é', 'e').replace('í', 'i').replace('ó', 'o').replace('ú', 'u')
               .replace('ñ', 'n'))
        if RX_TILDE.search(w) and wl not in LISTA_BLANCA_TILDE:
            out.append(('tilde/eñe «%s»' % w, ctx))
        elif wl in FUNCIONALES_ES and (wl != 'y' or w == 'y'):      # «Y» mayúscula = Yes (DV Y/N/Pending)
            out.append(('funcional ES «%s»' % w, ctx))
        elif sin in OFICIO_ES and wl not in LISTA_BLANCA_TILDE:
            out.append(('oficio ES «%s»' % w, ctx))
    m = re.search(r'(?<![\d,])\d+,\d{1,2}(?![\d])', t)
    if m:
        out.append(('coma decimal «%s»' % m.group(0), t[:80]))
    m = re.search(r'\b\d{2}:\d{2} ?- ?\d{2}:\d{2}\b', t)
    if m:
        out.append(('hora de 24 h «%s»' % m.group(0), t[:80]))
    m = RX_NORMA_ES.search(t)
    if m:
        out.append(('norma/sigla ES «%s»' % m.group(0), t[:80]))
    return out


# ==========================================================================
# Utilidades
# ==========================================================================
class Resultado:
    def __init__(self):
        self.gates = OrderedDict()
        self.actual = None

    def gate(self, nombre):
        self.actual = nombre
        self.gates.setdefault(nombre, {'fallos': [], 'datos': OrderedDict()})
        return self

    def fallo(self, msg):
        self.gates[self.actual]['fallos'].append(msg)

    def dato(self, k, v):
        self.gates[self.actual]['datos'][k] = v


def cargar(path, data_only=False):
    return openpyxl.load_workbook(path, data_only=data_only)


def libros(carpeta):
    for es in mapas.LIBROS_ES:
        en = mapas.FICHEROS[es]
        yield es, en, A.extraer_textos.corto(es), os.path.join(ORIGEN, es), os.path.join(carpeta, en)


def lits(f):
    return [x[1:-1].replace('""', '"') for x in RX_STR.findall(f or '')]


def celdas(sqref):
    out = []
    for rg in str(sqref).split():
        out.extend(A.rango(rg))
    return out


def es_verde(c):
    return (c.fill is not None and c.fill.fill_type == 'solid' and c.fill.fgColor is not None
            and isinstance(c.fill.fgColor.rgb, str) and c.fill.fgColor.rgb.upper().endswith(A.VERDE))


def _prot(ws):
    p = ws.protection
    return tuple(getattr(p, k) for k in ('sheet', 'password', 'sort', 'autoFilter', 'formatColumns', 'formatRows',
                                         'insertRows', 'deleteRows', 'selectLockedCells', 'selectUnlockedCells',
                                         'objects', 'scenarios', 'formatCells', 'insertColumns', 'deleteColumns'))


def _pagina(ws):
    ps = ws.page_setup
    fit = ws.sheet_properties.pageSetUpPr.fitToPage if ws.sheet_properties.pageSetUpPr else None
    return (ps.orientation, ps.fitToWidth, ps.fitToHeight, fit, str(ws.print_title_rows), str(ws.print_title_cols))


def _area(ws):
    pa = ws.print_area
    return pa.split('!')[-1] if pa else None


def _cf_firma(cf, regla, trans, trans_text):
    dxf = regla.dxf
    color = None
    if dxf is not None:
        color = (str(dxf.fill.fgColor.rgb) if dxf.fill is not None and dxf.fill.fgColor is not None else None,
                 str(dxf.font.color.rgb) if dxf.font is not None and dxf.font.color is not None else None)
    return (str(cf.sqref), regla.type, regla.operator, regla.priority, bool(regla.stopIfTrue),
            tuple(trans(x) for x in (regla.formula or [])), trans_text(regla.text), color)


def fechas_hoy(datos):
    out = defaultdict(dict)
    for x in datos['mercado']['fechas_hoy']:
        out[x['libro']][(x['hoja_en'], x['celda'])] = x['formula']
    return out


def dv_lista_en_esperada(f1_es):
    if f1_es == A.UNIDADES_ES_DV:
        return A.UNIDADES_EN_DV
    return mapas.dv_lista(A.CLAVES[x] for x in f1_es.strip('"').split(','))


# ==========================================================================
# G1 — paridad estructural
# ==========================================================================
def gateG1(R, carpeta, datos):
    R.gate('G1 paridad estructural')
    censo = datos['censo']
    d25 = fechas_hoy(datos)
    tot = Counter()
    for es, en, corto, pes, pen in libros(carpeta):
        if not os.path.isfile(pen):
            R.fallo('%s: no existe' % en)
            continue
        ce = censo['libros'][es]
        wes, wen = cargar(pes), cargar(pen)
        esperado = [mapas.HOJAS[h] for h in ce['hojas']]
        if wen.sheetnames != esperado or wes.sheetnames != ce['hojas']:
            R.fallo('%s: hojas %r ≠ %r' % (en, wen.sheetnames, esperado))
            continue
        mapa = OrderedDict((h, mapas.HOJAS[h]) for h in wes.sheetnames)
        T = A.Transformador(mapa, A.mapa_literales(corto))
        Tcf = A.Transformador(mapa, A.mapa_literales(corto, cf=True))
        for h_es in wes.sheetnames:
            ws_es, ws_en = wes[h_es], wen[mapa[h_es]]
            hd = ce['hojas_detalle'][h_es]
            lugar = '%s:%s' % (en, ws_en.title)
            f_es = {c.coordinate: c.value for c in ws_es._cells.values() if c.data_type == 'f'}
            f_en = {c.coordinate: c.value for c in ws_en._cells.values() if c.data_type == 'f'}
            dec = {k[1]: f for k, f in d25.get(corto, {}).items() if k[0] == ws_en.title}
            if len(f_es) != hd['celdas_formula']:
                R.fallo('%s: el ES ya no cuadra con el censo (%d ≠ %d fórmulas)' % (lugar, len(f_es), hd['celdas_formula']))
            if set(f_en) != set(f_es) | set(dec) or set(f_es) & set(dec):
                R.fallo('%s: celdas con fórmula distintas: solo ES %s · solo EN %s'
                        % (lugar, sorted(set(f_es) - set(f_en))[:5], sorted(set(f_en) - set(f_es) - set(dec))[:5]))
            for k, f in dec.items():
                tot['formulas_d25'] += 1
                if f_en.get(k) != f:
                    R.fallo('%s!%s: D25 %r ≠ %r' % (lugar, k, f_en.get(k), f))
            con_hoja = 0
            for k, fe in f_es.items():
                fn = f_en.get(k)
                if fn is None:
                    continue
                tot['formulas'] += 1
                try:
                    esperado_f = T.celda(fe)
                except A.Aborta as e:
                    esperado_f = '<<%s>>' % e
                if esperado_f != fn:
                    R.fallo('%s!%s: fórmula EN %r ≠ ES transformada %r' % (lugar, k, fn[:80], esperado_f[:80]))
                if A.extraer_textos.refs_hoja(fn):
                    con_hoja += 1
                    for h in A.extraer_textos.refs_hoja(fn):
                        if h not in wen.sheetnames:
                            R.fallo('%s!%s cita la hoja inexistente %r' % (lugar, k, h))
            tot['formulas_con_hoja'] += con_hoja
            if con_hoja != hd['formulas_con_hoja']:
                R.fallo('%s: %d fórmulas con referencia a hoja ≠ censo %d' % (lugar, con_hoja, hd['formulas_con_hoja']))
            # DV
            dves = sorted(ws_es.data_validations.dataValidation, key=lambda d: str(d.sqref))
            dven = sorted(ws_en.data_validations.dataValidation, key=lambda d: str(d.sqref))
            if len(dves) != len(dven) or len(dven) != len(hd['dv']):
                R.fallo('%s: %d DV ES, %d EN, %d censo' % (lugar, len(dves), len(dven), len(hd['dv'])))
            for de, dn in zip(dves, dven):
                tot['dv'] += 1
                tot['dv_celdas'] += len(celdas(dn.sqref))
                d11 = corto == '03' and str(de.sqref) == A.D11_DV_SQREF
                firma = lambda d: (str(d.sqref), bool(d.allow_blank), bool(d.showDropDown),  # noqa: E731
                                   bool(d.showErrorMessage), bool(d.showInputMessage), d.errorStyle)
                if firma(de) != firma(dn):
                    R.fallo('%s: DV %s cambia (%r → %r)' % (lugar, de.sqref, firma(de), firma(dn)))
                for attr in ('error', 'errorTitle', 'prompt', 'promptTitle'):
                    if bool(getattr(de, attr)) != bool(getattr(dn, attr)):
                        R.fallo('%s: DV %s pierde/gana %s' % (lugar, dn.sqref, attr))
                f1e, f1n = de.formula1 or '', dn.formula1 or ''
                if d11:
                    tot['dv_d11'] += 1
                    if (dn.type, dn.operator, f1n, dn.formula2) != ('decimal', 'between', '0', '100'):
                        R.fallo('%s: DV D11 %s = %r' % (lugar, dn.sqref, (dn.type, dn.operator, f1n, dn.formula2)))
                    continue
                if de.type != dn.type or de.operator != dn.operator:
                    R.fallo('%s: DV %s tipo %s/%s → %s/%s' % (lugar, dn.sqref, de.type, de.operator, dn.type, dn.operator))
                if f1e.startswith('"'):
                    try:
                        esp = dv_lista_en_esperada(f1e)
                    except KeyError as e:
                        esp = '<<sin clave %s>>' % e
                    if f1n != esp:
                        R.fallo('%s: DV %s lista %r ≠ %r' % (lugar, dn.sqref, f1n[:80], esp[:80]))
                    tot['dv_lista_literal'] += 1
                elif T.refs(f1e) != f1n:
                    R.fallo('%s: DV %s %r ≠ %r' % (lugar, dn.sqref, T.refs(f1e), f1n))
                else:
                    tot['dv_cita_hoja'] += 1
                if (de.formula2 or dn.formula2) and T.refs(de.formula2 or '') != (dn.formula2 or ''):
                    R.fallo('%s: DV %s formula2 %r ≠ %r' % (lugar, dn.sqref, de.formula2, dn.formula2))
            # CF
            tr_text = lambda t: Tcf.lit.get(t, t) if t else t       # noqa: E731
            cfe = sorted((_cf_firma(cf, r, Tcf.formula, tr_text) for cf in ws_es.conditional_formatting
                          for r in cf.rules), key=repr)
            cfn = sorted((_cf_firma(cf, r, lambda x: x, lambda x: x) for cf in ws_en.conditional_formatting
                          for r in cf.rules), key=repr)
            tot['cf'] += len(cfn)
            if len(cfn) != len(hd['cf']):
                R.fallo('%s: %d CF EN ≠ censo %d' % (lugar, len(cfn), len(hd['cf'])))
            if cfe != cfn:
                dif = [(a, b) for a, b in zip(cfe, cfn) if a != b] or [(cfe[:1], cfn[:1])]
                R.fallo('%s: formato condicional distinto (ES transformado → EN): %r' % (lugar, dif[0]))
            # merges, protección, paneles, impresión, autofiltro, columnas ocultas
            mg = sorted(map(str, ws_en.merged_cells.ranges))
            if mg != sorted(hd['merges']):
                R.fallo('%s: combinaciones distintas' % lugar)
            tot['merges'] += len(mg)
            if _prot(ws_es) != _prot(ws_en):
                R.fallo('%s: protección distinta' % lugar)
            if ws_en.protection.sheet:
                tot['hojas_protegidas'] += 1
                if ws_en.protection.password:
                    R.fallo('%s: protegida CON contraseña' % lugar)
            if ws_en.freeze_panes != hd['paneles']:
                R.fallo('%s: paneles %s ≠ %s' % (lugar, ws_en.freeze_panes, hd['paneles']))
            if _pagina(ws_es) != _pagina(ws_en):
                R.fallo('%s: ajuste o títulos de impresión %r ≠ %r' % (lugar, _pagina(ws_es), _pagina(ws_en)))
            if ws_en.print_title_rows or ws_en.print_title_cols:
                tot['titulos_impresion'] += 1
            if str(ws_en.auto_filter.ref or '') != str(hd['autofiltro'] or ''):
                R.fallo('%s: autofiltro %s ≠ %s' % (lugar, ws_en.auto_filter.ref, hd['autofiltro']))
            tot['autofiltros'] += 1 if ws_en.auto_filter.ref else 0
            ae = hd['area_impresion'].split('!')[-1] if hd['area_impresion'] else None
            if _area(ws_en) != ae:
                R.fallo('%s: área de impresión %s ≠ %s' % (lugar, _area(ws_en), ae))
            tot['areas_impresion'] += 1 if ae else 0
            ocultas = sorted(k for k, d in ws_en.column_dimensions.items() if d.hidden)
            if ocultas != hd['columnas_ocultas']:
                R.fallo('%s: columnas ocultas %s ≠ %s' % (lugar, ocultas, hd['columnas_ocultas']))
            tot['columnas_ocultas'] += len(ocultas)
            # verdes = desbloqueadas (+3 de D11 en Current Order)
            ue = {c.coordinate for c in ws_es._cells.values() if c.protection.locked is False}
            un = {c.coordinate for c in ws_en._cells.values() if c.protection.locked is False}
            extra = {c for c, _ in A.D11_DESGLOSE} if (corto == '03' and ws_en.title == 'Current Order') else set()
            if un != ue | extra:
                R.fallo('%s: celdas desbloqueadas distintas (solo ES %s, solo EN %s)'
                        % (lugar, sorted(ue - un)[:5], sorted(un - ue - extra)[:5]))
            verdes_desbl = [c for c in ws_en._cells.values() if c.protection.locked is False and es_verde(c)]
            if len(verdes_desbl) != len(un):
                R.fallo('%s: %d desbloqueadas y solo %d verdes' % (lugar, len(un), len(verdes_desbl)))
            verdes_bloq = [c.coordinate for c in ws_en._cells.values() if es_verde(c) and c.protection.locked is not False]
            if verdes_bloq:
                R.fallo('%s: celdas verdes bloqueadas %s' % (lugar, verdes_bloq[:5]))
            tot['desbloqueadas'] += len(un)
            tot['verdes_desbloqueadas'] += len(verdes_desbl)
        with zipfile.ZipFile(pen) as z:
            n = z.namelist()
        g = [x for x in n if x.startswith('xl/charts/') or x.startswith('xl/media/') or x.startswith('xl/drawings/')]
        if g:
            R.fallo('%s: gráficos/imágenes %s' % (en, g[:3]))
    T0 = censo['totales']
    esperados = OrderedDict([
        ('formulas', T0['celdas_formula']), ('formulas_d25', sum(len(v) for v in d25.values())),
        ('formulas_con_hoja', T0['formulas_con_hoja']), ('dv', T0['dv']), ('dv_celdas', T0['dv_celdas']),
        ('cf', T0['cf']), ('merges', T0['merges']), ('hojas_protegidas', T0['hojas_protegidas']),
        ('areas_impresion', T0['areas_impresion']), ('titulos_impresion', T0['titulos_impresion']),
        ('autofiltros', T0['autofiltros']), ('columnas_ocultas', T0['columnas_ocultas']),
        ('desbloqueadas', T0['desbloqueadas'] + len(A.D11_DESGLOSE)),
        ('verdes_desbloqueadas', T0['verdes_desbloqueadas'] + len(A.D11_DESGLOSE)),
    ])
    for k, v in esperados.items():
        if tot[k] != v:
            R.fallo('total %s = %d ≠ esperado %d (censo)' % (k, tot[k], v))
    if tot['dv_lista_literal'] + tot['dv_d11'] != T0['dv_lista_literal'] or tot['dv_cita_hoja'] != T0['dv_cita_hoja']:
        R.fallo('DV literales/citas %d+%d/%d ≠ censo %d/%d' % (tot['dv_lista_literal'], tot['dv_d11'], tot['dv_cita_hoja'],
                                                             T0['dv_lista_literal'], T0['dv_cita_hoja']))
    for k, v in sorted(tot.items()):
        R.dato(k, v)


# ==========================================================================
# G2 — restos y caracteres
# ==========================================================================
def piezas_de_texto(pen):
    wb = cargar(pen)
    for ws in wb.worksheets:
        yield 'hoja', ws.title, ws.title
        for c in ws._cells.values():
            if isinstance(c.value, str) and c.data_type != 'f':
                yield 'valor', '%s!%s' % (ws.title, c.coordinate), c.value
            if c.data_type == 'f':
                for x in lits(c.value):
                    yield 'literal', '%s!%s' % (ws.title, c.coordinate), x
            if isinstance(c.number_format, str) and c.number_format != 'General':
                yield 'formato', '%s!%s' % (ws.title, c.coordinate), c.number_format
        for dv in ws.data_validations.dataValidation:
            for attr in ('error', 'errorTitle', 'prompt', 'promptTitle', 'formula1'):
                v = getattr(dv, attr)
                if v and (attr != 'formula1' or v.startswith('"')):
                    yield 'dv', '%s DV %s %s' % (ws.title, dv.sqref, attr), v
        for cf in ws.conditional_formatting:
            for regla in cf.rules:
                if regla.text:
                    yield 'literal', '%s CF %s text' % (ws.title, cf.sqref), regla.text
                for f in regla.formula or []:
                    for x in lits(f):
                        yield 'literal', '%s CF %s' % (ws.title, cf.sqref), x
        for parte in ('oddHeader', 'oddFooter', 'evenHeader', 'evenFooter', 'firstHeader', 'firstFooter'):
            hf = getattr(ws, parte)
            for pos in ('left', 'center', 'right'):
                t = getattr(hf, pos).text
                if t:
                    yield 'cabecera-pie', '%s %s.%s' % (ws.title, parte, pos), t
    p = wb.properties
    for k in ('title', 'subject', 'creator', 'keywords', 'description', 'category', 'lastModifiedBy'):
        v = getattr(p, k)
        if v:
            yield 'docprops', k, v
    with zipfile.ZipFile(pen) as z:
        estilos = z.read('xl/styles.xml').decode('utf-8')
    for m in re.finditer(r'<numFmt [^>]*formatCode="([^"]*)"', estilos):
        yield 'styles.xml', 'numFmt', html.unescape(m.group(1))


def gateG2(R, carpeta, datos):
    R.gate('G2 restos')
    n = Counter()
    for es, en, corto, pes, pen in libros(carpeta):
        for tipo, lugar, t in piezas_de_texto(pen):
            n[tipo] += 1
            donde = '%s %s' % (en, lugar)
            if '€' in t or re.search(r'\bEUR\b', t):
                R.fallo('%s: moneda € en %s: %r' % (donde, tipo, t[:70]))
            if '$' in t:
                if tipo != 'valor':
                    R.fallo('%s: «$» en %s: %r' % (donde, tipo, t[:60]))
                elif len(re.findall(r'\$\d', t)) != t.count('$'):
                    R.fallo('%s: «$» que no es un importe: %r' % (donde, t[:80]))
                else:
                    n['importes_usd_en_texto'] += 1
            malos = no_permitidos(t)
            if malos:
                R.fallo('%s: carácter fuera de la lista blanca %r en %s: %r' % (donde, ''.join(malos), tipo, t[:60]))
            if tipo in ('formato', 'styles.xml'):
                if re.search(r'dd/mm', t, re.I):
                    R.fallo('%s: formato dd/mm %r' % (donde, t))
                continue
            m = RX_ID_INTERNO.search(t)
            if m:
                R.fallo('%s: ID interno %r en %s: …%s…' % (donde, m.group(0), tipo, t[max(0, m.start() - 30):m.end() + 10]))
            for motivo, ctx in restos_espanol(t):
                R.fallo('%s: %s · …%s…' % (donde, motivo, ctx))
    for k, v in sorted(n.items()):
        R.dato(k, v)


# ==========================================================================
# G3 — formato
# ==========================================================================
def gateG3(R, carpeta, datos, mes):
    R.gate('G3 formato')
    censo = datos['censo']
    d25 = fechas_hoy(datos)
    n = Counter()
    rx_version = re.compile(r'^Version %s · %s \d{4} · %s · info@aichef\.pro$'
                            % (re.escape(A.VERSION), mes, re.escape(A.DESCRIPTION)))
    for es, en, corto, pes, pen in libros(carpeta):
        ce = censo['libros'][es]
        wb = cargar(pen)
        for h_es, hd in ce['hojas_detalle'].items():
            ws = wb[mapas.HOJAS[h_es]]
            n['hojas'] += 1
            if ws.page_setup.paperSize not in (1, '1'):
                R.fallo('%s:%s: paperSize %s ≠ 1 (Letter)' % (en, ws.title, ws.page_setup.paperSize))
            for f in hd['fechas']:
                c = ws[f['celda']]
                n['fechas_censo'] += 1
                if c._style.numFmtId != 14:
                    R.fallo('%s:%s!%s: fecha con numFmtId %s ≠ 14' % (en, ws.title, f['celda'], c._style.numFmtId))
            for (hoja, celda) in d25.get(corto, {}):
                if hoja == ws.title:
                    n['fechas_d25'] += 1
                    if ws[celda]._style.numFmtId != 14:
                        R.fallo('%s:%s!%s: D25 sin numFmtId 14' % (en, ws.title, celda))
            fechas_es = {f['celda'] for f in hd['fechas']}
            for c in ws._cells.values():
                if c._style.numFmtId == 14 and c.coordinate not in fechas_es:
                    R.fallo('%s:%s!%s: numFmtId 14 en una celda que no era fecha' % (en, ws.title, c.coordinate))
                if isinstance(c.number_format, str) and re.search(r'dd/mm|€', c.number_format):
                    R.fallo('%s:%s!%s: formato %r' % (en, ws.title, c.coordinate, c.number_format))
            pies = {k: getattr(getattr(ws, k.split('.')[0]), k.split('.')[1]).text for k in hd['cabeceras_pies']}
            for k, v in pies.items():
                if v != A.PIE_EN:
                    R.fallo('%s:%s: %s = %r ≠ pie D19' % (en, ws.title, k, v))
                else:
                    n['pies'] += 1
        ins = wb[A.HOJA_INSTR_EN]
        vers = [c.value for c in ins._cells.values() if isinstance(c.value, str) and c.value.startswith('Version ')]
        if len(vers) != 1 or not rx_version.match(vers[0]):
            R.fallo('%s: línea de versión D21 %r' % (en, vers))
        else:
            n['versiones'] += 1
        p = wb.properties
        esp = {'title': mapas.TITULOS[en], 'subject': A.SUBJECT, 'keywords': A.KEYWORDS, 'description': A.DESCRIPTION,
               'category': A.CATEGORY, 'creator': A.CREATOR}
        for k, v in esp.items():
            if getattr(p, k) != v:
                R.fallo('%s: docProps %s %r ≠ %r' % (en, k, getattr(p, k), v))
        if ins['A1'].value != mapas.TITULOS[en]:
            R.fallo('%s: Instructions!A1 %r ≠ título D5' % (en, ins['A1'].value))
    T0 = censo['totales']
    if n['hojas'] != T0['hojas']:
        R.fallo('hojas %d ≠ censo %d' % (n['hojas'], T0['hojas']))
    if n['fechas_censo'] != T0['fechas']:
        R.fallo('fechas %d ≠ censo %d' % (n['fechas_censo'], T0['fechas']))
    for k, v in sorted(n.items()):
        R.dato(k, v)


# ==========================================================================
# G4 — claves
# ==========================================================================
def gateG4(R, carpeta, datos):
    R.gate('G4 claves')
    censo = datos['censo']
    cats = list(mapas.CATEGORIAS.values())
    lista_cats = mapas.dv_lista(cats)
    n = Counter()
    for es, en, corto, pes, pen in libros(carpeta):
        ce = censo['libros'][es]
        wes, wen = cargar(pes), cargar(pen)
        for h_es, hd in ce['hojas_detalle'].items():
            ws, ws_es = wen[mapas.HOJAS[h_es]], wes[h_es]
            lugar = '%s:%s' % (en, ws.title)
            tiene_uni = False
            for dv in ws.data_validations.dataValidation:
                f1 = dv.formula1 or ''
                if not f1.startswith('"'):
                    continue
                items = f1.strip('"').split(',')
                if f1 == lista_cats:
                    n['dv_categorias'] += 1
                elif any(x in cats for x in items):
                    R.fallo('%s: DV %s mezcla categorías: %r' % (lugar, dv.sqref, f1[:80]))
                if f1 == A.UNIDADES_EN_DV:
                    tiene_uni = True
                # todo valor de ejemplo dentro de su lista
                for coord in celdas(dv.sqref):
                    v = ws[coord].value
                    if v is None or (isinstance(v, str) and ws[coord].data_type == 'f'):
                        continue
                    n['valores_con_dv'] += 1
                    if str(v) not in items:
                        R.fallo('%s!%s: %r fuera de su DV %s' % (lugar, coord, v, f1[:60]))
            if tiene_uni != bool(hd['dv_unidades']):
                R.fallo('%s: lista de unidades D9 %s (ES %s)' % (lugar, tiene_uni, hd['dv_unidades']))
            n['hojas_dv_unidades'] += tiene_uni
            # etiquetas de agregación y ejemplos: toda celda que en ES era una categoría es una categoría EN
            for c in ws_es._cells.values():
                if isinstance(c.value, str) and c.value in mapas.CATEGORIAS and c.data_type != 'f' and \
                        not (ws.title == 'Receiving Temps' and c.column_letter in 'AH'):   # familias («Congelados»)
                    v = ws[c.coordinate].value
                    n['celdas_categoria'] += 1
                    if v not in cats:
                        R.fallo('%s!%s: %r no es una de las 10 categorías' % (lugar, c.coordinate, v))
            # fórmulas SUMIF/COUNTIF por categoría: los literales de categoría (si hubiera) son EN
            for c in ws._cells.values():
                if c.data_type == 'f':
                    for x in lits(c.value):
                        if x in mapas.CATEGORIAS:
                            R.fallo('%s!%s: literal de categoría ES %r' % (lugar, c.coordinate, x))
    if n['hojas_dv_unidades'] != censo['totales']['hojas_dv_unidades']:
        R.fallo('hojas con DV de unidades %d ≠ censo %d' % (n['hojas_dv_unidades'], censo['totales']['hojas_dv_unidades']))
    # familias (04) y zonas (06)
    w4 = cargar(os.path.join(carpeta, mapas.FICHEROS['04-recepcion-mercancias.xlsx']))['Receiving Temps']
    for i, (en_f, mn, mx, ideal, base) in enumerate(mapas.FAMILIAS.values()):
        r = 4 + i
        got = (w4['A%d' % r].value, w4['B%d' % r].value, w4['C%d' % r].value, w4['D%d' % r].value)
        if got != (en_f, ideal, mn, mx):
            R.fallo('04 Receiving Temps fila %d: %r ≠ %r' % (r, got, (en_f, ideal, mn, mx)))
        n['familias'] += 1
    for i, (cat, fam) in enumerate(mapas.FAMILIA_POR_CATEGORIA.items()):
        r = 4 + i
        if (w4['G%d' % r].value, w4['H%d' % r].value) != (cat, fam):
            R.fallo('04 Receiving Temps G/H%d: %r' % (r, (w4['G%d' % r].value, w4['H%d' % r].value)))
    w6 = cargar(os.path.join(carpeta, mapas.FICHEROS['06-fifo-caducidades.xlsx']))
    zonas = list(mapas.ZONAS.values())
    got = [w6['Storage Map']['A%d' % r].value for r in range(4, 4 + len(zonas))]
    if got != zonas:
        R.fallo('06 Storage Map A4:A14 %r ≠ mapas.ZONAS' % got)
    dvz = [dv for dv in w6['FIFO Tracker'].data_validations.dataValidation if 'R5' in celdas(dv.sqref)]
    if not dvz or dvz[0].formula1 != mapas.dv_lista(zonas):
        R.fallo('06 DV de zonas ≠ mapas.ZONAS')
    n['zonas'] = len(zonas)
    for k, v in sorted(n.items()):
        R.dato(k, v)


# ==========================================================================
# G5 — CF y literales
# ==========================================================================
def gateG5(R, carpeta, datos):
    R.gate('G5 CF y literales')
    n = Counter()
    for es, en, corto, pes, pen in libros(carpeta):
        wb = cargar(pen)
        todos = set()
        dv_items = set()
        for ws in wb.worksheets:
            for c in ws._cells.values():
                if c.data_type == 'f':
                    todos.update(lits(c.value))
            for dv in ws.data_validations.dataValidation:
                if (dv.formula1 or '').startswith('"'):
                    dv_items.update(dv.formula1.strip('"').split(','))
        for ws in wb.worksheets:
            lugar = '%s:%s' % (en, ws.title)
            for cf in ws.conditional_formatting:
                estados = set()
                for coord in celdas(cf.sqref):
                    c = ws[coord]
                    if c.data_type == 'f':
                        estados.update(x for x in lits(c.value) if x and not x.startswith('*'))
                for regla in cf.rules:
                    if regla.type == 'containsText':
                        tok = regla.text
                        casan = sorted(e for e in estados if tok.lower() in e.lower())
                        n['tokens_cf'] += 1
                        ok = len(casan) == 1 or (tok == 'REJECT' and len(casan) == 2 and
                                                  all(e.startswith('✗ REJECT') for e in casan))
                        if not ok:
                            R.fallo('%s: token CF %r casa con %r (rango %s)' % (lugar, tok, casan, cf.sqref))
                        if 'SEARCH("%s"' % tok not in ''.join(regla.formula or []):
                            R.fallo('%s: CF %s text %r ≠ su fórmula %r' % (lugar, cf.sqref, tok, regla.formula))
                    else:
                        for f in regla.formula or []:
                            for x in lits(f):
                                if not x:
                                    continue
                                n['literales_cf_expresion'] += 1
                                if not any(x.lower() in t.lower() for t in todos if t != x) and x not in dv_items:
                                    R.fallo('%s: CF %s busca %r y ninguna fórmula lo devuelve' % (lugar, cf.sqref, x))
            for c in ws._cells.values():
                if c.data_type != 'f':
                    continue
                for x in lits(c.value):
                    if len(x) > 2 and x.startswith('*') and x.endswith('*'):
                        n['comodines'] += 1
                        nucleo = x.strip('*')
                        if not any(nucleo in t for t in todos if not t.startswith('*')):
                            R.fallo('%s!%s: comodín %r sin estado que lo produzca' % (lugar, c.coordinate, x))
    # literales comparados con su DV
    w2 = cargar(os.path.join(carpeta, mapas.FICHEROS['02-fichas-proveedores.xlsx']))
    aprob = [dv.formula1 for dv in w2['Vendor Directory'].data_validations.dataValidation
             if dv.formula1 == mapas.dv_lista(mapas.DV_LISTAS['homologado'].values())]
    usa_y = any('"Y"' in (c.value or '') for c in w2['Vendor Scorecard']._cells.values() if c.data_type == 'f')
    if not aprob or not usa_y:
        R.fallo('02: "Y" del Scorecard sin su DV Y/N/Pending (DV %s, literal %s)' % (bool(aprob), usa_y))
    w6 = cargar(os.path.join(carpeta, mapas.FICHEROS['06-fifo-caducidades.xlsx']))
    ft = [dv.formula1 for dv in w6['FIFO Tracker'].data_validations.dataValidation if 'F5' in celdas(dv.sqref)]
    usa_bb = any('"Best-by"' in (c.value or '') for c in w6['FIFO Tracker']._cells.values() if c.data_type == 'f')
    if ft != [mapas.dv_lista(mapas.DV_LISTAS['tipo_fecha'].values())] or not usa_bb:
        R.fallo('06: "Best-by" sin su DV Use-by/Best-by (%r, %s)' % (ft, usa_bb))
    w7 = cargar(os.path.join(carpeta, mapas.FICHEROS['07-analisis-costes-compras.xlsx']))
    per = [dv.formula1 for dv in w7['KPI Dashboard'].data_validations.dataValidation if 'B10' in celdas(dv.sqref)]
    esp = mapas.dv_lista([mapas.PERIODO_ANIO[1]] + list(mapas.MESES_CORTOS.values()))
    usa_fy = any('"Full year"' in (c.value or '') for c in w7['KPI Dashboard']._cells.values() if c.data_type == 'f')
    cab = [w7['Spend by Category'].cell(3, c).value for c in range(2, 14)]
    if per != [esp] or not usa_fy or cab != list(mapas.MESES_CORTOS.values()):
        R.fallo('07: periodo DV %r / "Full year" %s / cabeceras %r' % (per, usa_fy, cab))
    n['literales_con_dv'] = 3
    for k, v in sorted(n.items()):
        R.dato(k, v)


# ==========================================================================
# G6 — cálculo
# ==========================================================================
def celdas_sin_v(path):
    fuera = set()
    with zipfile.ZipFile(path) as z:
        wbx = z.read('xl/workbook.xml').decode('utf-8')
        rels = z.read('xl/_rels/workbook.xml.rels').decode('utf-8')
        destino = {}
        for tag in re.findall(r'<Relationship\b[^>]*>', rels):
            i, t = re.search(r'\bId="([^"]+)"', tag), re.search(r'\bTarget="([^"]+)"', tag)
            if i and t:
                destino[i.group(1)] = t.group(1)
        for tag in re.findall(r'<sheet\b[^>]*>', wbx):
            nombre = html.unescape(re.search(r'\bname="([^"]*)"', tag).group(1))
            rid = re.search(r'\br:id="([^"]+)"', tag).group(1)
            t = destino[rid]
            parte = t.lstrip('/') if t.startswith('/') else 'xl/' + t
            xml = z.read(parte).decode('utf-8')
            for m in re.finditer(r'<c\b([^>]*?)(?:/>|>(.*?)</c>)', xml, re.S):
                dentro = m.group(2) or ''
                if '<f' in dentro and not re.search(r'<v>[^<]', dentro) and '<v></v>' not in dentro \
                        and '<v/>' not in dentro and '<v />' not in dentro:
                    fuera.add((nombre, re.search(r'\br="([A-Z]+\d+)"', m.group(1)).group(1)))
                elif '<f' in dentro and ('<v/>' in dentro or '<v />' in dentro):
                    fuera.add((nombre, re.search(r'\br="([A-Z]+\d+)"', m.group(1)).group(1)))
    return fuera


def num(v):
    return isinstance(v, (int, float)) and not isinstance(v, bool)


def gateG6(R, carpeta, datos):
    R.gate('G6 cálculo')
    n = Counter()
    for es, en, corto, pes, pen in libros(carpeta):
        wf, wv = cargar(pen), cargar(pen, data_only=True)
        sin_v = celdas_sin_v(pen)
        for ws in wf.worksheets:
            for c in ws._cells.values():
                if c.data_type != 'f':
                    continue
                n['formulas'] += 1
                if (ws.title, c.coordinate) in sin_v:
                    R.fallo('%s:%s!%s sin caché (<f> sin <v>)' % (en, ws.title, c.coordinate))
                    continue
                v = wv[ws.title][c.coordinate].value
                if v is None:
                    n['vacias_legitimas'] += 1
                elif isinstance(v, str) and v in ERRORES:
                    R.fallo('%s:%s!%s = %s' % (en, ws.title, c.coordinate, v))
                else:
                    n['con_valor'] += 1
    M = datos['mercado']['identidades']
    # historia del 04
    w4 = cargar(os.path.join(carpeta, mapas.FICHEROS['04-recepcion-mercancias.xlsx']), True)['Receiving Log']
    for x in M['recepcion_04']:
        v = w4['R%d' % x['fila']].value
        if v != x['estado_esperado']:
            R.fallo('04 Receiving Log!R%d = %r ≠ %r' % (x['fila'], v, x['estado_esperado']))
    hist = [w4['R%d' % x['fila']].value for x in M['recepcion_04']]
    if hist != ['✓ ACCEPT', '✗ REJECT (too warm)', '✓ ACCEPT', 'N/A']:
        R.fallo('04 historia %r ≠ ACCEPT / REJECT (too warm) / ACCEPT / N/A (SPEC G6)' % hist)
    n['historia_04'] = len(hist)
    # 07 food cost dentro del objetivo
    w7 = cargar(os.path.join(carpeta, mapas.FICHEROS['07-analisis-costes-compras.xlsx']), True)['KPI Dashboard']
    fc, obj, est = w7['B4'].value, w7['C4'].value, w7['E4'].value
    if not (num(fc) and num(obj) and fc <= obj and 0.28 <= fc <= 0.32 and abs(fc - M['food_cost_07']) < 0.0005
            and est == '\U0001f7e2 OK'):
        R.fallo('07 food cost %r (objetivo %r, estado %r, identidad %r)' % (fc, obj, est, M['food_cost_07']))
    R.dato('food_cost_07', fc)
    # 06: los 5 estados hoy
    w6 = cargar(os.path.join(carpeta, mapas.FICHEROS['06-fifo-caducidades.xlsx']), True)['FIFO Tracker']
    estados = [w6['L%d' % x['fila']].value for x in M['fifo_06']]
    for x, v in zip(M['fifo_06'], estados):
        if v != x['estado_esperado']:
            R.fallo('06 FIFO Tracker!L%d = %r ≠ %r' % (x['fila'], v, x['estado_esperado']))
    base = {e.split(' ', 1)[1] if ' ' in e else e for e in estados if e}
    if len({mapas.LITERALES.get(k, k) for k in estados}) < 5:
        R.fallo('06: %d estados distintos hoy (%r), se prometen 5' % (len(set(estados)), estados))
    R.dato('estados_06', sorted(set(estados)))
    for k, v in sorted(n.items()):
        R.dato(k, v)


# ==========================================================================
# G7 — identidades D14
# ==========================================================================
def gateG7(R, carpeta, datos):
    R.gate('G7 identidades')
    r = subprocess.run([sys.executable, VALIDAR_MERCADO], stdout=subprocess.PIPE, stderr=subprocess.STDOUT,
                       universal_newlines=True)
    ult = r.stdout.strip().splitlines()[-1] if r.stdout.strip() else ''
    R.dato('validar_mercado', ult)
    if r.returncode:
        R.fallo('validar_mercado.py exit %d: %s' % (r.returncode, ult))
    M = datos['mercado']
    n = Counter()
    wbs = {}
    for es, en, corto, pes, pen in libros(carpeta):
        wbs[corto] = (cargar(pen), cargar(pen, data_only=True))
    for corto, lst in M['libros'].items():
        wf = wbs[corto][0]
        for e in lst:
            v = wf[e['hoja_en']][e['celda']].value
            esp = e['valor']
            if isinstance(esp, str) and A.RX_FECHA_ISO.match(esp):
                esp = datetime.datetime.strptime(esp, '%Y-%m-%d')
            if isinstance(v, datetime.datetime) and isinstance(esp, datetime.datetime):
                igual = v == esp
            elif num(v) and num(esp):
                igual = abs(v - esp) < 1e-9
            else:
                igual = v == esp
            n['celdas_mercado'] += 1
            if not igual:
                R.fallo('%s %s!%s = %r ≠ mercado %r' % (corto, e['hoja_en'], e['celda'], v, esp))
    # mismo producto ⇒ mismo precio en la hoja (01 y BONUS-08 contra la tabla maestra)
    prod = {p['en']: p for p in M['productos']}
    w1 = wbs['01'][0]
    for p in M['productos']:
        l1 = p.get('libro01')
        if not l1:
            continue
        ws = w1[l1['hoja_en']]
        r = l1['fila']
        got = (ws['B%d' % r].value, ws['C%d' % r].value, ws['D%d' % r].value, ws['E%d' % r].value,
               ws['F%d' % r].value, ws['J%d' % r].value)
        esp = (p['en'], p['categoria_en'], p['unidad'], p['par'], p['par_max'], p['precio'])
        if got != esp:
            R.fallo('01 %s fila %d: %r ≠ tabla maestra %r' % (l1['hoja_en'], r, got, esp))
        n['maestra_01'] += 1
    # BONUS-09: punto de pedido = par del 01; stock máx = par max (valores de la caché)
    wv9 = wbs['B09'][1]['Reorder Calculator']
    par01 = {p['id']: (p['par'], p['par_max']) for p in M['productos']}
    for x in M['identidades']['b09']:
        rr = x['fila']
        rop, mx, sug, eoq = (wv9['H%d' % rr].value, wv9['K%d' % rr].value, wv9['M%d' % rr].value,
                             wv9['L%d' % rr].value)
        par, pmax = par01[x['id']]
        if not (num(rop) and abs(rop - par) < TOL and mx == pmax):
            R.fallo('B09 fila %d (%s): ROP %r / máx %r ≠ par 01 %r / %r' % (rr, x['id'], rop, mx, par, pmax))
        if not (num(sug) and abs(sug - x['sugerida']) < TOL and num(eoq) and abs(eoq - x['eoq']) < TOL):
            R.fallo('B09 fila %d: sugerida/EOQ %r/%r ≠ %r/%r' % (rr, sug, eoq, x['sugerida'], x['eoq']))
        n['b09'] += 1
    I = M['identidades']
    # importes = cantidad × precio (caché)
    w3 = wbs['03'][1]['Current Order']
    if not (num(w3['H40'].value) and abs(w3['H40'].value - I['pedido_actual_03']['subtotal']) < TOL):
        R.fallo('03 subtotal %r ≠ %r' % (w3['H40'].value, I['pedido_actual_03']['subtotal']))
    for r in range(9, 39):
        e, f, g = w3['E%d' % r].value, w3['F%d' % r].value, w3['G%d' % r].value
        if num(e) and num(f):
            n['lineas_03'] += 1
            if not (num(g) and abs(g - e * f) < TOL):
                R.fallo('03 Current Order!G%d = %r ≠ %r × %r' % (r, g, e, f))
    if w3['A43'].value != I['pedido_actual_03']['estado']:
        R.fallo('03 A43 %r ≠ %r' % (w3['A43'].value, I['pedido_actual_03']['estado']))
    wl = wbs['03'][1]['Order Log']
    for h in I['historial_03']:
        fila = [wl.cell(h['fila'], c).value for c in range(1, 12)]
        nums = [x for x in fila if num(x)]
        if not any(abs(x - h['total']) < TOL for x in nums):
            R.fallo('03 Order Log fila %d sin el total %r: %r' % (h['fila'], h['total'], fila))
    w4 = wbs['04'][1]['Receiving Log']
    for x in I['recepcion_04']:
        v = w4['M%d' % x['fila']].value
        if x.get('valor_diferencia') is not None and not (num(v) and abs(abs(v) - abs(x['valor_diferencia'])) < TOL):
            R.fallo('04 Receiving Log!M%d = %r ≠ %r' % (x['fila'], v, x['valor_diferencia']))
    w5 = wbs['05'][1]
    g105 = w5['Daily Waste Log']['G105'].value
    if not (num(g105) and abs(g105 - I['mermas_05_total']) < TOL):
        R.fallo('05 total %r ≠ %r' % (g105, I['mermas_05_total']))
    for r in range(4, 104):
        ws = w5['Daily Waste Log']
        d, f, g = ws['D%d' % r].value, ws['F%d' % r].value, ws['G%d' % r].value
        if num(d) and num(f):
            n['lineas_05'] += 1
            if not (num(g) and abs(g - d * f) < TOL):
                R.fallo('05 Daily Waste Log!G%d = %r ≠ %r × %r' % (r, g, d, f))
    if w5['Waste Dashboard']['B10'].value != I['compras_mes_05']:
        R.fallo('05 compras %r ≠ %r' % (w5['Waste Dashboard']['B10'].value, I['compras_mes_05']))
    w7 = wbs['07'][1]
    b14 = w7['Spend by Category']['B14'].value
    if not (num(b14) and abs(b14 - I['compras_enero_07']) < TOL and b14 == w5['Waste Dashboard']['B10'].value):
        R.fallo('07 compras de enero %r ≠ %r / 05 %r' % (b14, I['compras_enero_07'], w5['Waste Dashboard']['B10'].value))
    kp = w7['KPI Dashboard']
    if not (num(kp['B5'].value) and abs(kp['B5'].value - I['coste_por_cubierto_07']) < 0.01):
        R.fallo('07 coste por cubierto %r ≠ %r' % (kp['B5'].value, I['coste_por_cubierto_07']))
    for x in I['top20_07']:
        ws = w7['Top 20 Items']
        g = [ws.cell(x['fila'], c).value for c in range(1, 14)]
        if not any(num(v) and abs(v - x['gasto_mes']) < TOL for v in g):
            R.fallo('07 Top 20 fila %d sin el gasto %r: %r' % (x['fila'], x['gasto_mes'], g))
    # cifras D20 recalculadas sobre la hoja EN final: toda cifra con $ o % del texto 07 A19 sale de la caché
    a19 = wbs['07'][0]['KPI Dashboard']['A19'].value
    fig = {'$%s' % format(int(round(v)), ',') for v in (kp['B6'].value, kp['B14'].value, kp['B15'].value,
                                                         kp['B12'].value) if num(v)}
    for f in fig:
        if f not in a19:
            R.fallo('07 KPI Dashboard!A19 no cita %s (caché): %r' % (f, a19[:120]))
    if '%.1f%%' % (kp['B4'].value * 100) not in a19 or '{:,}'.format(int(kp['B13'].value)) not in a19:
        R.fallo('07 A19 no cita el food cost %.1f%% o los cubiertos' % (kp['B4'].value * 100))
    for (libro, hoja, celda), en in A.cargar_datos()['derivados'].items():
        v = wbs[libro][0][hoja][celda].value
        n['d20'] += 1
        if v != en:
            R.fallo('D20 %s %s!%s no es la reescritura de mercado_en.json' % (libro, hoja, celda))
    w9p = wbs['B09'][0]['Parameters']
    pb = M['parametros_b09']
    vals = [w9p.cell(r, c).value for r in range(1, 20) for c in range(1, 6)]
    for k in ('coste_pedido', 'almacenamiento_anual', 'factor_vida_util'):
        if pb[k] not in vals:
            R.fallo('B09 Parameters sin %s = %r' % (k, pb[k]))
    for k, v in sorted(n.items()):
        R.dato(k, v)


# ==========================================================================
# G8 — censo-entregables
# ==========================================================================
def gateG8(R, carpeta, datos):
    R.gate('G8 censo-entregables')
    if os.path.abspath(carpeta) == os.path.abspath(A.DESTINO_REAL):
        cmd = [sys.executable, CENSO_SCRIPT, '--only', mapas.SLUG, '--fail']
    else:
        cmd = [sys.executable, CENSO_SCRIPT, '--only', carpeta, '--letter', '--fail']
    r = subprocess.run(cmd, stdout=subprocess.PIPE, stderr=subprocess.STDOUT, universal_newlines=True, cwd=REPO)
    lineas = [l for l in r.stdout.strip().splitlines() if l.strip()]
    R.dato('comando', ' '.join(cmd[1:]))
    R.dato('salida', lineas[-3:])
    if r.returncode:
        R.fallo('censo-entregables exit %d: %s' % (r.returncode, ' | '.join(lineas[-5:])))


# ==========================================================================
# Autotest: cada gate tiene que cazar un defecto inyectado en una COPIA
# ==========================================================================
def _mutar(carpeta, fname, fn):
    p = os.path.join(carpeta, fname)
    wb = cargar(p)
    fn(wb)
    wb.save(p)


def _cf_token(ws, de, a):
    for cf in ws.conditional_formatting:
        for r in cf.rules:
            if r.text == de:
                r.text = a
                r.formula = [f.replace('"%s"' % de, '"%s"' % a) for f in r.formula]


MUTACIONES = [
    ('G1', '01-kitchen-bar-inventory-par-sheet.xlsx',
     lambda wb: setattr(wb['Kitchen']['H5'], 'value', wb['Kitchen']['H5'].value.replace('REORDER', 'PEDIR')), 'fórmula EN'),
    ('G1', '03-purchase-order-template.xlsx',
     lambda wb: setattr(wb['Current Order']['A46'], 'protection', openpyxl.styles.Protection(locked=True)), 'desbloqueadas'),
    ('G2', '05-food-waste-log.xlsx', lambda wb: setattr(wb['Action Plan']['A1'], 'value', 'Plan de la merma'), 'ES'),
    ('G2', '02-vendor-list-price-comparison.xlsx',
     lambda wb: setattr(wb['Vendor Directory']['A1'], 'value', 'Vendors kitinv-a · list'), 'ID interno'),
    ('G3', '06-fifo-expiration-date-tracker.xlsx',
     lambda wb: setattr(wb['FIFO Tracker'].page_setup, 'paperSize', 9), 'paperSize'),
    ('G4', '07-purchasing-cost-analysis.xlsx',
     lambda wb: setattr(wb['Spend by Category']['A4'], 'value', 'Meat and Poultry'), 'categor'),
    ('G5', '01-kitchen-bar-inventory-par-sheet.xlsx', lambda wb: _cf_token(wb['Kitchen'], 'REORDER', 'O'), 'token CF'),
    ('G6', '04-receiving-log.xlsx', lambda wb: setattr(wb['Receiving Log']['O5'], 'value', 38), 'REJECT'),
    ('G7', 'BONUS-09-reorder-point-calculator.xlsx',
     lambda wb: setattr(wb['Reorder Calculator']['D4'], 'value', 10), 'B09'),
]


def autotest(carpeta, mes, datos):
    ok = True
    base = ejecutar(carpeta, mes, None, datos, silencioso=True)
    sucios = [g for g, v in base.gates.items() if v['fallos']]
    if sucios:
        print('AUTOTEST: la carpeta base no está limpia (%s): arregla eso primero' % sucios)
        return False
    for gate, fname, fn, aguja in MUTACIONES:
        tmp = tempfile.mkdtemp(prefix='rik-autotest-')
        try:
            for f in mapas.FICHEROS.values():
                shutil.copy2(os.path.join(carpeta, f), tmp)
            _mutar(tmp, fname, fn)
            if gate in ('G6', 'G7'):
                subprocess.run([sys.executable, '-W', 'ignore', A.INJECT, os.path.join(tmp, fname)],
                               stdout=subprocess.PIPE, stderr=subprocess.PIPE)
            R = ejecutar(tmp, mes, [gate], datos, silencioso=True)
            fallos = [f for g in R.gates.values() for f in g['fallos']]
            cazado = any(aguja.lower() in f.lower() for f in fallos)
            ok &= cazado
            print('  autotest %s %-40s %s' % (gate, fname, 'CAZADO' if cazado else 'NO CAZADO %r' % fallos[:2]))
        finally:
            shutil.rmtree(tmp)
    return ok


# ==========================================================================
# main
# ==========================================================================
GATES = OrderedDict([('G1', gateG1), ('G2', gateG2), ('G3', gateG3), ('G4', gateG4), ('G5', gateG5),
                     ('G6', gateG6), ('G7', gateG7), ('G8', gateG8)])


def cargar_datos():
    return {'censo': json.load(open(os.path.join(AQUI, 'censo_es.json'), encoding='utf-8')),
            'mercado': json.load(open(os.path.join(AQUI, 'mercado_en.json'), encoding='utf-8'))}


def ejecutar(carpeta, mes, solo, datos, silencioso=False):
    R = Resultado()
    for nombre, fn in GATES.items():
        if solo and nombre not in solo:
            continue
        A.termica()
        if nombre == 'G3':
            fn(R, carpeta, datos, mes)
        else:
            fn(R, carpeta, datos)
        g = R.gates[R.actual]
        if not silencioso:
            estado = 'OK' if not g['fallos'] else 'FALLA (%d)' % len(g['fallos'])
            resumen = ' · '.join('%s=%s' % (k, v) for k, v in list(g['datos'].items())[:8] if not isinstance(v, list))
            print('%-24s %s  %s' % (R.actual, estado, resumen[:200]))
            for f in g['fallos'][:int(os.environ.get('RIK_MAX_FALLOS', '12'))]:
                print('    - ' + f[:220])
    return R


def main():
    ap = argparse.ArgumentParser(description='Gates F2 del Restaurant Inventory Kit Pro (EN)')
    ap.add_argument('--dir', default=A.DESTINO_REAL, help='carpeta con los 9 xlsx EN (por defecto dl/<slug>)')
    ap.add_argument('--mes', default=None, help='mes de la línea de versión (por defecto, el actual)')
    ap.add_argument('--solo', default=None, help='G1,G2,… (por defecto, todos)')
    ap.add_argument('--json', default=None)
    ap.add_argument('--autotest', action='store_true')
    args = ap.parse_args()
    mes = args.mes or A.mes_por_defecto()
    datos = cargar_datos()
    carpeta = os.path.abspath(args.dir)
    if args.autotest:
        ok = autotest(carpeta, mes, datos)
        print('AUTOTEST: %s' % ('OK' if ok else 'FALLA'))
        sys.exit(0 if ok else 1)
    solo = [s.strip().upper() for s in args.solo.split(',')] if args.solo else None
    R = ejecutar(carpeta, mes, solo, datos)
    fallos = sum(len(g['fallos']) for g in R.gates.values())
    if args.json:
        with open(args.json, 'w', encoding='utf-8') as fh:
            json.dump(R.gates, fh, ensure_ascii=False, indent=1, default=str)
    print('\n%s: %d gates, %d fallos (%s)' % ('VERDE' if not fallos else 'ROJO', len(R.gates), fallos, carpeta))
    sys.exit(1 if fallos else 0)


if __name__ == '__main__':
    main()
