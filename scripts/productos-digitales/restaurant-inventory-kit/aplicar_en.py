#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
aplicar_en.py — Restaurant Inventory Kit Pro (EN) · montaje de los 9 xlsx desde el Kit de Inventario v2.0 publicado.

SPEC que manda: `SPEC.md` (§2 mapas, §4 temperaturas, §5 datos, §7 pasos 4-6, D5-D25). Sesión Claude Code,
25-sep-2026. Copia adaptada de `recipe-costing-kit/aplicar_en.py` (el piloto): reutiliza su maquinaria genérica
(guarda térmica, renombrado de hojas con reescritura de referencias, literales en fórmulas/DV/CF, traducción por
celda, formatos sin «€» y fechas a numFmtId 14, US Letter + pie D19, docProps D21, apóstrofos rectos) y añade las
funciones de mercado PROPIAS de este kit. Aquí no se redacta nada: los textos salen de `textos_en/` (subagentes
Anthropic, regla 1bis), `mercado_en.json` (tabla maestra, proveedores, fechas D25, textos D20) y `mapas.py`.

    python3 aplicar_en.py                  # escribe en astro-site/public/dl/restaurant-inventory-templates/
    python3 aplicar_en.py --dry-run        # escribe en $CLAUDE_SCRATCHPAD/rik-dryrun (o .work/<slug>/rik-dryrun)
    python3 aplicar_en.py --dry-run --salida <carpeta> [--mes September] [--idempotencia] [--json informe.json]

Los ficheros ES de `dl/kit-inventario/` se abren en SOLO LECTURA y nunca se escriben.

Orden por libro, siempre desde una copia limpia del ES publicado (→ idempotente):
   1. carga del xlsx ES;
   2. renombrado de hojas (mapas.HOJAS) reescribiendo TODAS las referencias (fórmulas, DV, áreas y títulos de
      impresión viajan con el título en openpyxl);
   3. literales de fórmula (mapas.LITERALES) en celdas y CF; tokens de CF (mapas.TOKENS_CF, con el ajuste del 06);
   4. textos por celda (textos_en + localizar + claves de mapas), cabeceras/pies, línea de versión D21;
      después, los textos con cifra derivada de mercado_en.json en su celda (D20: mandan sobre la traducción);
   5. DV: listas literales por clave (D8), unidades D9, D11 (lista «4,10,21» → decimal 0-100), mensajes EN sin
      IDs internos de auditoría;
   6. mercado: `libros` celda a celda (D14/D15), `fechas_hoy` como fórmulas TODAY()+k (D25), `vaciar` (D11);
   7. D11 (03: Tax Rates a 0 %, desglose 0/5/20 editable y verde), D12 (04: Receiving Temps reescrita por fila
      en °F con mapas.FAMILIAS; 06: Storage Map desde `libros`);
   8. formatos: sin «€» (D13), fechas a numFmtId 14, US Letter y pie D19;
   9. docProps D21, título interior D5 comprobado, apóstrofos rectos;
  10. guardado con el nombre EN (mapas.FICHEROS);
  11. inject_cache.py AL FINAL sobre los 9 (cualquier wb.save() posterior borraría la caché) y comprobación
      con load_workbook(data_only=True) de que las fórmulas de ejemplo muestran valores.

Térmica: todo en SERIE; `istats` antes de cada libro y de cada inject_cache (espera si ≥ 62 °C hasta < 60 °C).
"""
import argparse
import copy
import datetime
import glob
import json
import logging
import os
import re
import shutil
import subprocess
import sys
import time
from collections import OrderedDict

AQUI = os.path.dirname(os.path.abspath(__file__))
SCRIPTS = os.path.dirname(AQUI)
REPO = os.path.dirname(os.path.dirname(SCRIPTS))
sys.path.insert(0, AQUI)

import openpyxl                                              # noqa: E402
from openpyxl.styles import PatternFill, Protection         # noqa: E402

import mapas                                                 # noqa: E402
import extraer_textos                                        # noqa: E402

logging.disable(logging.CRITICAL)

ORIGEN = os.path.join(REPO, 'astro-site', 'public', 'dl', 'kit-inventario')
DESTINO_REAL = os.path.join(REPO, 'astro-site', 'public', 'dl', mapas.SLUG)   # restaurant-inventory-templates
INJECT = os.path.join(SCRIPTS, 'inject_cache.py')

VERSION = '2.0'
SUBJECT = mapas.PRODUCTO + ' · v2.0'                                  # D21
DESCRIPTION = mapas.URL_PRODUCTO                                        # aichef.pro/en/digital-products/…
KEYWORDS = 'restaurant inventory template, par sheet, AI Chef Pro'
CATEGORY = 'AI Chef Pro · Digital products'
CREATOR = 'AI Chef Pro'
PIE_EN = 'AI Chef Pro · aichef.pro · Page &P of &N'                   # D19
MESES = ['January', 'February', 'March', 'April', 'May', 'June', 'July', 'August',
         'September', 'October', 'November', 'December']
VERDE = 'E8F5E9'
FMT_FECHA = 'mm-dd-yy'                     # numFmtId 14 = fecha corta del sistema (D19)
INVARIABLES = {'OK', 'N/A', 'AI Chef Pro'}
RX_VERSION_ES = re.compile(r'^Versi[óo]n \d+\.\d+ · ')
RX_FECHA_ISO = re.compile(r'^\d{4}-\d{2}-\d{2}$')
# Prefijos de auditoría que el ES dejó en los títulos de mensaje de DV («grupo_b · …», «kitinv-v2 · …»): son IDs
# internos (como el «RT-08:» de I1) y no llegan al comprador EN (D23, mismo criterio que I1).
RX_ID_DV = re.compile(r'^(?:kitinv-[a-z0-9]+|grupo_[a-z0-9]+) · ')

# Textos SIN letras (fuera de textos_es.json por el extractor) que llevan convención española: «€», coma decimal u
# hora de 24 h. Se fijan aquí, en su forma US (D13, D19). Única fuente: estas 8 cadenas del ES publicado.
SIN_LETRAS = {
    '2-5 €': '2-5 (in your currency)',                 # B09 Parameters!B4 (D24: coste de pedido neutro)
    '0,7 (70 %)': '0.7 (70%)',                          # B09 Parameters!B12
    '07:00 - 09:00': '7:00-9:00 a.m.',                  # 02 Vendor Terms!G4:G9 (ventana de reparto)
    '06:30 - 08:00': '6:30-8:00 a.m.',
    '06:00 - 08:00': '6:00-8:00 a.m.',
    '09:00 - 13:00': '9:00 a.m.-1:00 p.m.',
    '10:00 - 13:00': '10:00 a.m.-1:00 p.m.',
    '11:00 - 14:00': '11:00 a.m.-2:00 p.m.',
}

# Cadenas ES con dos sentidos según la celda: mandan en su celda (libro, hoja EN, celda) → (ES esperado, EN).
# «Nota» es la NOTA A/B/C/D del proveedor en 02 (→ «Grade») y una columna de notas en 03 «Tax Rates».
POR_CELDA = {
    ('03', 'Tax Rates', 'C1'): ('Nota', 'Note'),
}

HOJA_INSTR_EN = mapas.HOJAS['Instrucciones']
UNIDADES_ES_DV = mapas.dv_lista(mapas.UNIDADES.keys())          # "kg,L,ud,docena,caja,bandeja,barril,saco,rollo,paquete"
UNIDADES_EN_DV = mapas.dv_lista(mapas.UNIDADES_EN)
CLAVES = OrderedDict((es, v[1]) for es, v in mapas.todas_las_claves().items())

# D11 · 03 «Current Order» / «Tax Rates»
D11_DV_SQREF = 'H9:H38'
D11_DV_ES = '"4,10,21"'
D11_DESGLOSE = (('A46', 0), ('A47', 5), ('A48', 20))
D11_CATEGORIAS = 'B2:B11'
D11_ERROR_TITULO = 'Invalid tax rate'
D11_ERROR = ('Type a tax rate between 0 and 100 (a number, without the % sign). Most food and drinks bought '
             'for resale are 0%; use the rate you actually pay.')
D11_NOTA_A49 = ('Breakdown by tax rate: type the rates you actually pay in the green cells of column A (0, 5 and '
                '20 are only a starting point; UK VAT uses 0, 5 and 20%). The sum of the Net amount column has to '
                'match the NET AMOUNT (PRE-TAX) above.')


class Aborta(Exception):
    pass


# ==========================================================================
# Térmica (regla del Mac: se apaga por encima de ~65 °C)
# ==========================================================================
def termica(umbral=62, reanuda=60):
    def leer():
        try:
            out = subprocess.run(['istats', 'cpu', 'temp', '--value-only'], stdout=subprocess.PIPE,
                                 stderr=subprocess.PIPE, universal_newlines=True, timeout=20).stdout
            return float(re.findall(r'[\d.]+', out)[0])
        except Exception:
            return None
    t = leer()
    if t is None or t < umbral:
        return t
    while t is not None and t >= reanuda:
        print('  [térmica] CPU %.1f °C: espero 30 s' % t, flush=True)
        time.sleep(30)
        t = leer()
    return t


# ==========================================================================
# Datos
# ==========================================================================
def cargar_datos():
    te = json.load(open(os.path.join(AQUI, 'textos_es.json'), encoding='utf-8'))
    tr = OrderedDict()

    def poner(es, en, origen):
        if es in tr and tr[es] != en:
            raise Aborta('traducción contradictoria para %r (%s): %r ≠ %r' % (es[:60], origen, tr[es][:60], en[:60]))
        tr[es] = en
    for p in sorted(glob.glob(os.path.join(AQUI, 'textos_en', '*.json'))):
        for x in json.load(open(p, encoding='utf-8')):
            poner(x['es'], x['en'], os.path.basename(p))
    mercado = json.load(open(os.path.join(AQUI, 'mercado_en.json'), encoding='utf-8'))
    for x in mercado['localizar']:
        poner(x['es'], x['en'], 'mercado_en.json localizar')
    for c in te['cadenas']:
        if c['tratamiento'] in ('mapa-clave', 'mapa-hoja'):
            poner(c['es'], c['en_mapa'], 'mapa')
    # D20: textos con cifra derivada → mandan en SU celda (libro corto, hoja EN, celda)
    derivados = {(x['book'], x['sheet'], x['cell']): x['en'] for x in mercado['textos_cifra_derivada']}

    def cubierta_por_d20(c):
        return all(d['t'] == 'celda' and (d['f'], mapas.HOJAS[d['h']], d['c']) in derivados for d in c['donde'])
    faltan = [c['es'] for c in te['cadenas'] if c['tratamiento'] != 'regenerar' and c['es'] not in tr
              and not cubierta_por_d20(c)]
    if faltan:
        raise Aborta('%d cadenas de textos_es.json sin traducción: %r' % (len(faltan), faltan[:5]))
    # D14/D15: celdas de mercado por libro
    libros = {}
    for corto, lst in mercado['libros'].items():
        d = OrderedDict()
        for e in lst:
            k = (e['hoja_en'], e['celda'])
            if k in d and d[k] != e['valor']:
                raise Aborta('mercado_en.json: %s %s!%s con dos valores' % (corto, e['hoja_en'], e['celda']))
            d[k] = e['valor']
        libros[corto] = d
    fechas = {}
    for x in mercado['fechas_hoy']:
        fechas.setdefault(x['libro'], OrderedDict())[(x['hoja_en'], x['celda'])] = x
    vaciar = {}
    for x in mercado['vaciar']:
        vaciar.setdefault(x['libro'], []).append(x)
    return {'te': te, 'tr': tr, 'mercado': mercado, 'derivados': derivados, 'libros': libros,
            'fechas': fechas, 'vaciar': vaciar}


def mes_por_defecto():
    return MESES[datetime.date.today().month - 1]


def linea_version(mes):
    return 'Version %s · %s %d · %s · info@aichef.pro' % (VERSION, mes, datetime.date.today().year, DESCRIPTION)


# ==========================================================================
# Fórmulas: hojas y literales
# ==========================================================================
RX_STR = re.compile(r'"(?:[^"]|"")*"')
RX_REF = re.compile(r"(?:'((?:[^']|'')+)'|([^\W\d][\w.]*))!")


def ref_hoja(nombre):
    """Nombre de hoja tal como va en una fórmula (entre comillas si hace falta)."""
    if re.fullmatch(r'[A-Za-z_][A-Za-z0-9_.]*', nombre):
        return nombre
    return "'" + nombre.replace("'", "''") + "'"


def mapa_literales(corto, cf=False):
    """Literales ES→EN de un libro. En CF entran además los tokens de containsText (con el ajuste del 06)."""
    m = OrderedDict(mapas.LITERALES)
    if cf:
        tok = OrderedDict(mapas.TOKENS_CF)
        tok.update(mapas.TOKENS_CF_POR_LIBRO.get(corto, {}))
        for k, v in tok.items():
            m.setdefault(k, v)
    return m


class Transformador:
    """Reescribe una fórmula ES a EN: referencias de hoja (fuera de literales) y literales (dentro)."""

    def __init__(self, mapa_hojas, literales):
        self.mapa = mapa_hojas
        self.en = set(mapa_hojas.values())
        self.lit = literales
        self.hojas_citadas = 0
        self.literales = 0

    def refs(self, s):
        def rep(m):
            nombre = (m.group(1) if m.group(1) is not None else m.group(2)).replace("''", "'")
            if nombre in self.mapa:
                self.hojas_citadas += 1
                return ref_hoja(self.mapa[nombre]) + '!'
            if nombre in self.en:
                return m.group(0)
            raise Aborta('hoja desconocida en fórmula: %r en %r' % (nombre, s[:80]))
        return RX_REF.sub(rep, s)

    def formula(self, f):
        out, pos = [], 0
        for m in RX_STR.finditer(f):
            out.append(self.refs(f[pos:m.start()]))
            lit = m.group(0)[1:-1].replace('""', '"')
            if lit in self.lit:
                lit = self.lit[lit]
                self.literales += 1
            out.append('"' + lit.replace('"', '""') + '"')
            pos = m.end()
        out.append(self.refs(f[pos:]))
        return ''.join(out)

    def celda(self, v):
        return '=' + self.formula(v[1:]) if v.startswith('=') else self.formula(v)


def formato_en(fmt):
    """number_format EN: sin «€» (D13) y fechas a la corta del sistema (D19)."""
    if not isinstance(fmt, str):
        return fmt
    if fmt.lower() in ('dd/mm/yyyy', 'dd/mm/yy', 'd/m/yyyy'):
        return FMT_FECHA
    if '€' in fmt:
        f = re.sub(r'\[\$€[^\]]*\]', '', fmt)
        f = f.replace('"€"', '').replace('€', '')
        return re.sub(r'\s+', ' ', f).strip()
    return fmt


# ==========================================================================
# Escritura protegida (el mercado nunca pisa una fórmula)
# ==========================================================================
CONTADOR = {'mercado': 0}


def escribir(ws, coord, valor):
    c = ws[coord]
    if c.data_type == 'f' or (isinstance(c.value, str) and c.value.startswith('=')):
        raise Aborta('%s!%s es una fórmula (%r): el mercado no la pisa' % (ws.title, coord, c.value))
    if isinstance(valor, str) and RX_FECHA_ISO.match(valor):
        valor = datetime.datetime.strptime(valor, '%Y-%m-%d')
    c.value = valor
    CONTADOR['mercado'] += 1


def rango(ref):
    c1, r1, c2, r2 = openpyxl.utils.cell.range_boundaries(ref)
    for r in range(r1, r2 + 1):
        for c in range(c1, c2 + 1):
            yield '%s%d' % (openpyxl.utils.get_column_letter(c), r)


# ==========================================================================
# Pasos genéricos (del piloto)
# ==========================================================================
def renombrar(wb):
    mapa = OrderedDict((ws.title, mapas.HOJAS[ws.title]) for ws in wb.worksheets)
    for ws in wb.worksheets:
        ws.title = mapa[ws.title]
    if wb.sheetnames != list(mapa.values()):
        raise Aborta('renombrado: %r ≠ %r' % (wb.sheetnames, list(mapa.values())))
    return mapa


def reescribir_formulas(wb, T, rep):
    n = 0
    for ws in wb.worksheets:
        for c in list(ws._cells.values()):
            if c.data_type == 'f' and isinstance(c.value, str):
                nuevo = T.celda(c.value)
                if nuevo != c.value:
                    c.value = nuevo
                    n += 1
    rep['formulas_reescritas'] = n


def celdas_regeneradas(corto):
    """Celdas que NO se traducen por cadena: las reescribe una función de mercado (regenerar, D11, D12)."""
    s = set()
    if corto == '04':
        s |= {('Receiving Temps', c) for c in rango('A4:E18')}
        s |= {('Receiving Temps', c) for c in rango('G4:H13')}
    if corto == '03':
        s |= {('Tax Rates', c) for x in [None] for c in rango('E2:F20')}
    return s


def traducir_libro(wb, corto, fname_en, mapa, datos, mes, rep):
    """Textos de celda (no fórmula), cabeceras/pies y línea de versión; luego los textos D20 en su celda."""
    tr = datos['tr']
    faltan = rep.setdefault('faltan', [])
    mercado = datos['libros'].get(corto, {})
    fijas = celdas_regeneradas(corto)
    n = 0

    def traducir(v, lugar):
        if v in tr:
            return tr[v]
        if v in SIN_LETRAS:
            return SIN_LETRAS[v]
        if v in INVARIABLES or not any(ch.isalpha() for ch in v):
            return v
        faltan.append('%s: %r' % (lugar, v[:90]))
        return v

    for ws in wb.worksheets:
        for c in list(ws._cells.values()):
            v = c.value
            if not isinstance(v, str) or c.data_type == 'f':
                continue
            k = (ws.title, c.coordinate)
            if k in mercado or k in fijas:
                continue                                   # lo escribe el mercado
            if RX_VERSION_ES.match(v):
                c.value = linea_version(mes)
                n += 1
                continue
            if (corto, ws.title, c.coordinate) in datos['derivados']:
                continue                                   # D20: se escribe abajo
            pc = POR_CELDA.get((corto, ws.title, c.coordinate))
            if pc:
                if v != pc[0]:
                    raise Aborta('POR_CELDA %s %s!%s: el ES trae %r y no %r' % (corto, ws.title, c.coordinate, v, pc[0]))
                c.value = pc[1]
                n += 1
                continue
            nuevo = traducir(v, '%s!%s' % (ws.title, c.coordinate))
            if nuevo != v:
                c.value = nuevo
                n += 1
        for parte in ('oddHeader', 'oddFooter', 'evenHeader', 'evenFooter', 'firstHeader', 'firstFooter'):
            hf = getattr(ws, parte)
            for pos in ('left', 'center', 'right'):
                t = getattr(hf, pos).text
                if not t:
                    continue
                if 'Página &P' in t:
                    getattr(hf, pos).text = PIE_EN                   # D19 (regenerar)
                else:
                    getattr(hf, pos).text = traducir(t, '%s %s.%s' % (ws.title, parte, pos))
    # D20: cifra derivada recalculada (mercado_en.json) — manda sobre la traducción literal
    d20 = 0
    for (libro, hoja, celda), en in datos['derivados'].items():
        if libro != corto:
            continue
        c = wb[hoja][celda]
        if not isinstance(c.value, str) or c.data_type == 'f':
            raise Aborta('D20 %s %s!%s no es una celda de texto: %r' % (corto, hoja, celda, c.value))
        c.value = en
        d20 += 1
    rep['textos_traducidos'] = n
    rep['textos_d20'] = d20
    return traducir


def reescribir_dv_cf(wb, corto, T, Tcf, traducir, rep):
    dv_uni = dv_lista = dv_ref = cf_n = 0
    for ws in wb.worksheets:
        for dv in ws.data_validations.dataValidation:
            f1 = dv.formula1 or ''
            if f1 == UNIDADES_ES_DV:
                dv.formula1 = UNIDADES_EN_DV                        # D9: se sustituye, no se mapea
                dv_uni += 1
            elif corto == '03' and str(dv.sqref) == D11_DV_SQREF and f1 == D11_DV_ES:
                dv.type = 'decimal'                                   # D11: tipo de impuesto libre 0-100
                dv.operator = 'between'
                dv.formula1 = '0'
                dv.formula2 = '100'
                rep['d11_dv'] = True
            elif f1.startswith('"'):
                items = f1.strip('"').split(',')
                faltan = [x for x in items if x not in CLAVES]
                if faltan:
                    raise Aborta('%s: DV %s con ítems sin clave EN %r' % (ws.title, dv.sqref, faltan))
                dv.formula1 = mapas.dv_lista(CLAVES[x] for x in items)
                if len(dv.formula1) > 257:
                    raise Aborta('%s: DV %s > 255 caracteres' % (ws.title, dv.sqref))
                dv_lista += 1
            elif f1:
                dv.formula1 = T.refs(f1)
                dv_ref += 1
            if dv.formula2 and not (corto == '03' and str(dv.sqref) == D11_DV_SQREF):
                dv.formula2 = T.refs(dv.formula2)
            for attr in ('error', 'errorTitle', 'prompt', 'promptTitle'):
                v = getattr(dv, attr)
                if v:
                    nuevo = RX_ID_DV.sub('', traducir(v, '%s DV %s %s' % (ws.title, dv.sqref, attr)))
                    setattr(dv, attr, nuevo)
            if corto == '03' and str(dv.sqref) == D11_DV_SQREF:
                dv.errorTitle = D11_ERROR_TITULO
                dv.error = D11_ERROR
        for cf in ws.conditional_formatting:
            for regla in cf.rules:
                if regla.formula:
                    regla.formula = [Tcf.formula(x) for x in regla.formula]
                    cf_n += 1
                if regla.text:
                    lit = Tcf.lit
                    if regla.text not in lit:
                        raise Aborta('%s: CF %s con texto sin token EN %r' % (ws.title, cf.sqref, regla.text))
                    regla.text = lit[regla.text]
    rep.update(dv_unidades=dv_uni, dv_listas=dv_lista, dv_ref=dv_ref, cf_reglas=cf_n)


def formatos_y_papel(wb, rep):
    """D13 (sin moneda), D19 (numFmtId 14) y US Letter."""
    nf = wb._number_formats
    for i, f in enumerate(list(nf)):
        if '€' in f:
            list.__setitem__(nf, i, formato_en(f))
    _indice_formatos(nf)
    fechas = cambiados = 0
    for ws in wb.worksheets:
        for c in list(ws._cells.values()):
            f = c.number_format
            if isinstance(f, str) and (f == FMT_FECHA or f.lower() in ('dd/mm/yyyy', 'dd/mm/yy')):
                c.number_format = FMT_FECHA            # builtin 14
                fechas += 1
            else:
                g = formato_en(f)
                if g != f:
                    c.number_format = g
                    cambiados += 1
        ws.page_setup.paperSize = ws.PAPERSIZE_LETTER
    for i, f in enumerate(list(nf)):
        if f.lower() in ('dd/mm/yyyy', 'dd/mm/yy'):
            libre = 'm/d/yyyy'
            while libre in list(nf):
                libre += ';@'
            list.__setitem__(nf, i, libre)
    _indice_formatos(nf)
    rep.update(celdas_fecha=fechas, formatos_cambiados=cambiados)


def _indice_formatos(nf):
    if len(set(nf)) != len(nf):
        raise Aborta('formatos propios repetidos: %r' % list(nf))
    nf._dict = {v: i for i, v in enumerate(nf)}
    nf.clean = True


APOSTROFOS = {'\u2019': "'", '\u2018': "'"}


def normalizar_apostrofos(wb, rep):
    n = 0
    for ws in wb.worksheets:
        for c in list(ws._cells.values()):
            if isinstance(c.value, str) and c.data_type != 'f' and any(a in c.value for a in APOSTROFOS):
                v = c.value
                for a, b in APOSTROFOS.items():
                    v = v.replace(a, b)
                c.value = v
                n += 1
        for dv in ws.data_validations.dataValidation:
            for attr in ('error', 'errorTitle', 'prompt', 'promptTitle'):
                v = getattr(dv, attr)
                if v and any(a in v for a in APOSTROFOS):
                    for a, b in APOSTROFOS.items():
                        v = v.replace(a, b)
                    setattr(dv, attr, v)
                    n += 1
    rep['apostrofos_normalizados'] = n


def metadatos(wb, fname_en, rep):
    p = wb.properties
    p.title = mapas.TITULOS[fname_en]                 # D5
    p.subject = SUBJECT
    p.keywords = KEYWORDS
    p.description = DESCRIPTION
    p.category = CATEGORY
    p.creator = CREATOR
    p.lastModifiedBy = CREATOR
    rep['docprops'] = OrderedDict((k, getattr(p, k)) for k in
                                  ('title', 'subject', 'creator', 'keywords', 'description', 'category'))


def comprobar_titulo(wb, fname_en, rep):
    """D5: el título interior (Instructions!A1) tiene que ser EXACTAMENTE el de mapas.TITULOS. Nunca se parchea."""
    esperado = mapas.TITULOS[fname_en]
    v = wb[HOJA_INSTR_EN]['A1'].value
    if v != esperado:
        raise Aborta('%s: Instructions!A1 %r ≠ D5 %r' % (fname_en, v, esperado))
    rep['titulo'] = esperado


# ==========================================================================
# Mercado propio del kit
# ==========================================================================
def mercado_libros(wb, corto, datos, rep):
    """D14/D15: las celdas de ejemplo de mercado_en.json, una a una (mandan sobre la sustitución por cadena)."""
    celdas = datos['libros'].get(corto, {})
    for (hoja, celda), valor in celdas.items():
        escribir(wb[hoja], celda, valor)
    rep['celdas_mercado'] = len(celdas)


def mercado_fechas_hoy(wb, corto, datos, rep):
    """D25: fechas de ejemplo vivas, =TODAY()+k (mismo desfase que el ES respecto a su fecha de referencia)."""
    n = 0
    for (hoja, celda), x in datos['fechas'].get(corto, {}).items():
        c = wb[hoja][celda]
        if c.data_type == 'f':
            raise Aborta('D25 %s!%s ya es una fórmula' % (hoja, celda))
        if not isinstance(c.value, datetime.datetime) or c.value.strftime('%Y-%m-%d') != x['fecha_es']:
            raise Aborta('D25 %s!%s: el ES trae %r y no %s' % (hoja, celda, c.value, x['fecha_es']))
        if not re.fullmatch(r'=TODAY\(\)[+-]\d+', x['formula']):
            raise Aborta('D25 %s!%s: fórmula inesperada %r' % (hoja, celda, x['formula']))
        c.value = x['formula']
        c.number_format = FMT_FECHA
        n += 1
    rep['fechas_today'] = n


def mercado_vaciar(wb, corto, datos, rep):
    n = 0
    for x in datos['vaciar'].get(corto, []):
        ws = wb[x['hoja_en']]
        for coord in rango(x['rango']):
            if ws[coord].value is not None:
                escribir(ws, coord, None)
                n += 1
    rep['celdas_vaciadas'] = n


def mercado_03(wb, rep):
    """D11: impuesto neutro. Tax Rates al 0 % en las 10 categorías; desglose 0/5/20 editable y verde (corrige I7)."""
    ws = wb['Tax Rates']
    for coord in rango(D11_CATEGORIAS):
        escribir(ws, coord, 0)
    # la tabla de excepciones queda vacía (vaciar) con su cabecera; las filas 11-20 sin estilo heredan el de la 10
    ws = wb['Current Order']
    verde = None
    for c in ws._cells.values():
        if c.protection.locked is False and c.fill is not None and c.fill.fill_type == 'solid' and \
                str(c.fill.fgColor.rgb).upper().endswith(VERDE):
            verde = c
            break
    if verde is None:
        raise Aborta('03: no encuentro una celda verde de referencia en Current Order')
    for coord, valor in D11_DESGLOSE:
        c = ws[coord]
        if c.data_type == 'f':
            raise Aborta('03 %s es fórmula' % coord)
        c.value = valor
        c.fill = PatternFill('solid', fgColor=VERDE)
        c.protection = Protection(locked=False)
    a49 = ws['A49']
    if not isinstance(a49.value, str) or not a49.value.startswith('Breakdown by tax rate'):
        raise Aborta('03 A49 inesperada: %r' % a49.value)
    a49.value = D11_NOTA_A49
    rep['d11'] = 'Tax Rates 0 %% ×10 · desglose %s editable' % '/'.join(str(v) for _, v in D11_DESGLOSE)


def mercado_04(wb, rep):
    """D12: Receiving Temps reescrita POR FILA con mapas.FAMILIAS (°F, FDA Food Code 2022) y el puente
    categoría → familia (mapas.FAMILIA_POR_CATEGORIA). El rango de la DV y de los VLOOKUP no cambia."""
    ws = wb['Receiving Temps']
    for i, (es, (en, mn, mx, ideal, base)) in enumerate(mapas.FAMILIAS.items()):
        r = 4 + i
        if ws['A%d' % r].value != es:
            raise Aborta('04 Receiving Temps!A%d = %r, esperaba %r' % (r, ws['A%d' % r].value, es))
        escribir(ws, 'A%d' % r, en)
        escribir(ws, 'B%d' % r, ideal)
        escribir(ws, 'C%d' % r, mn)
        escribir(ws, 'D%d' % r, mx)
        escribir(ws, 'E%d' % r, limpiar_etiquetas(base))
    cats = list(mapas.CATEGORIAS.items())
    for i, (cat_es, cat_en) in enumerate(cats):
        r = 4 + i
        if ws['G%d' % r].value != cat_es:
            raise Aborta('04 Receiving Temps!G%d = %r, esperaba %r' % (r, ws['G%d' % r].value, cat_es))
        escribir(ws, 'G%d' % r, cat_en)
        escribir(ws, 'H%d' % r, mapas.FAMILIA_POR_CATEGORIA[cat_en])
    rep['d12_familias'] = len(mapas.FAMILIAS)


RX_ETIQUETA = re.compile(r'\s*\[(?:fuente|estimado|derivado)\]')


def limpiar_etiquetas(t):
    """Las etiquetas de trazabilidad de la SPEC ([fuente]/[estimado]) y las notas internas en español no llegan al
    comprador: «Base» queda con la cita normativa en inglés."""
    t = RX_ETIQUETA.sub('', t)
    t = t.replace('; max = límite físico', '; max = physical limit')
    t = t.replace('«received frozen»', '"received frozen"')
    t = re.sub(r';\s*min$', '', t.strip())
    t = re.sub(r';\s*0 °F$', '', t.strip())
    t = re.sub(r';\s*range$', '', t.strip())
    return t.strip().rstrip(';').strip()


def aplicar_mercado(wb, corto, datos, rep):
    CONTADOR['mercado'] = 0
    mercado_vaciar(wb, corto, datos, rep)
    mercado_libros(wb, corto, datos, rep)
    mercado_fechas_hoy(wb, corto, datos, rep)
    if corto == '03':
        mercado_03(wb, rep)
    if corto == '04':
        mercado_04(wb, rep)
    rep['escrituras_mercado'] = CONTADOR['mercado']


# ==========================================================================
# Un libro
# ==========================================================================
def construir_libro(fname_es, carpeta, datos, mes):
    fname_en = mapas.FICHEROS[fname_es]
    corto = extraer_textos.corto(fname_es)
    src = os.path.join(ORIGEN, fname_es)
    dst = os.path.join(carpeta, fname_en)
    rep = OrderedDict([('es', fname_es), ('en', fname_en)])
    wb = openpyxl.load_workbook(src)

    mapa = renombrar(wb)                                                 # 2
    T = Transformador(mapa, mapa_literales(corto))
    Tcf = Transformador(mapa, mapa_literales(corto, cf=True))
    reescribir_formulas(wb, T, rep)                                      # 2-3
    traducir = traducir_libro(wb, corto, fname_en, mapa, datos, mes, rep)   # 4
    reescribir_dv_cf(wb, corto, T, Tcf, traducir, rep)                   # 3, 5
    rep['refs_hoja'] = T.hojas_citadas + Tcf.hojas_citadas
    rep['literales'] = T.literales + Tcf.literales
    aplicar_mercado(wb, corto, datos, rep)                               # 6-7
    formatos_y_papel(wb, rep)                                            # 8
    metadatos(wb, fname_en, rep)                                         # 9
    comprobar_titulo(wb, fname_en, rep)
    normalizar_apostrofos(wb, rep)
    if rep['faltan']:
        raise Aborta('%s: %d textos sin traducción:\n  %s'
                     % (fname_es, len(rep['faltan']), '\n  '.join(rep['faltan'][:20])))
    wb.save(dst)                                                         # 10
    return rep


def construir(carpeta, datos, mes, log=print):
    os.makedirs(carpeta, exist_ok=True)
    informe = []
    for fname_es in mapas.LIBROS_ES:
        termica()
        rep = construir_libro(fname_es, carpeta, datos, mes)
        informe.append(rep)
        log('  %-44s fórmulas %4d · refs %4d · literales %4d · textos %4d · D20 %d · mercado %4d · fechas %3d'
            % (rep['en'], rep['formulas_reescritas'], rep['refs_hoja'], rep['literales'],
               rep['textos_traducidos'], rep['textos_d20'], rep['escrituras_mercado'], rep['celdas_fecha']))
    return informe


# ==========================================================================
# Idempotencia (huella comparable, sin fechas de guardado)
# ==========================================================================
def digest(path):
    wb = openpyxl.load_workbook(path)
    p = wb.properties
    fuera = {'_props': (p.title, p.subject, p.keywords, p.description, p.category, p.creator)}
    for ws in wb.worksheets:
        celdas = {}
        for c in ws._cells.values():
            if c.value is None:
                continue
            relleno = str(c.fill.fgColor.rgb) if c.fill is not None and c.fill.fill_type == 'solid' else None
            celdas[c.coordinate] = (repr(c.value), c.number_format, relleno, bool(c.protection.locked))
        fuera[ws.title] = {
            'celdas': celdas,
            'merges': sorted(str(m) for m in ws.merged_cells.ranges),
            'dv': sorted('%s:%s:%s:%s' % (dv.type, dv.formula1, dv.formula2, dv.sqref)
                         for dv in ws.data_validations.dataValidation),
            'cf': sorted('%s:%s' % (cf.sqref, [r.formula for r in cf.rules]) for cf in ws.conditional_formatting),
            'prot': bool(ws.protection.sheet), 'area': str(ws.print_area),
            'papel': ws.page_setup.paperSize, 'pie': ws.oddFooter.center.text,
        }
    return fuera


def diferencias(a, b, fichero):
    out = []
    if a['_props'] != b['_props']:
        out.append('%s: propiedades' % fichero)
    for hoja in sorted((set(a) | set(b)) - {'_props'}):
        if hoja not in a or hoja not in b:
            out.append('%s:%s: hoja solo en una pasada' % (fichero, hoja))
            continue
        for k in ('merges', 'dv', 'cf', 'prot', 'area', 'papel', 'pie'):
            if a[hoja][k] != b[hoja][k]:
                out.append('%s:%s: cambia %s' % (fichero, hoja, k))
        ca, cb = a[hoja]['celdas'], b[hoja]['celdas']
        for coord in sorted(set(ca) | set(cb)):
            if ca.get(coord) != cb.get(coord):
                out.append('%s:%s!%s: %s → %s' % (fichero, hoja, coord, ca.get(coord), cb.get(coord)))
    return out


# ==========================================================================
# inject_cache y verificación de la caché
# ==========================================================================
def inject_cache(carpeta, log=print):
    fuera = []
    for fname_es in mapas.LIBROS_ES:
        termica()
        p = os.path.join(carpeta, mapas.FICHEROS[fname_es])
        r = subprocess.run([sys.executable, '-W', 'ignore', INJECT, p], stdout=subprocess.PIPE,
                           stderr=subprocess.PIPE, universal_newlines=True)
        linea = r.stdout.strip().splitlines()[-1] if r.stdout.strip() else ''
        log('    ' + linea)
        fuera.append({'fichero': mapas.FICHEROS[fname_es], 'salida': linea, 'exit': r.returncode,
                      'stderr': r.stderr[-400:] if r.returncode else ''})
    return fuera


def verificar_cache(carpeta):
    """Toda fórmula de una fila con datos de ejemplo tiene valor en data_only (None solo si el resultado es "")."""
    out = OrderedDict()
    for fname_es in mapas.LIBROS_ES:
        p = os.path.join(carpeta, mapas.FICHEROS[fname_es])
        wf, wv = openpyxl.load_workbook(p), openpyxl.load_workbook(p, data_only=True)
        n = con_valor = 0
        for ws in wf.worksheets:
            for c in ws._cells.values():
                if c.data_type == 'f':
                    n += 1
                    if wv[ws.title][c.coordinate].value is not None:
                        con_valor += 1
        out[mapas.FICHEROS[fname_es]] = (n, con_valor)
    return out


def carpeta_dryrun(arg):
    if arg:
        return os.path.abspath(arg)
    s = os.environ.get('CLAUDE_SCRATCHPAD')
    if s:
        return os.path.join(os.path.abspath(s), 'rik-dryrun')
    return os.path.join(REPO, '.work', mapas.SLUG, 'rik-dryrun')


def main():
    ap = argparse.ArgumentParser(description='Restaurant Inventory Kit Pro (EN): montaje de los 9 xlsx')
    ap.add_argument('--dry-run', action='store_true', help='escribe en el scratchpad (o --salida), no en dl/')
    ap.add_argument('--salida', default=None, help='carpeta del dry-run')
    ap.add_argument('--mes', default=None, help='mes de la línea de versión D21 (por defecto, el actual)')
    ap.add_argument('--json', default=None, help='informe JSON')
    ap.add_argument('--idempotencia', action='store_true', help='2.ª construcción en un clon y 0 diferencias')
    ap.add_argument('--sin-cache', action='store_true', help='(depuración) no corre inject_cache')
    args = ap.parse_args()

    mes = args.mes or mes_por_defecto()
    if mes not in MESES:
        raise SystemExit('--mes debe ser un mes en inglés: %s' % ', '.join(MESES))
    carpeta = carpeta_dryrun(args.salida) if args.dry_run else DESTINO_REAL
    if os.path.abspath(carpeta) == os.path.abspath(ORIGEN):
        raise SystemExit('ABORTADO: la salida no puede ser la carpeta ES')

    t0 = time.time()
    try:
        datos = cargar_datos()
        print('== 1/3 · montaje de los 9 libros → %s (mes %s)' % (carpeta, mes), flush=True)
        informe = construir(carpeta, datos, mes)
    except Aborta as e:
        raise SystemExit('ABORTADO: %s' % e)
    # nada más en la carpeta que los 9 EN (idempotente también en ficheros)
    sobran = sorted(set(os.listdir(carpeta)) - set(mapas.FICHEROS.values()))
    if sobran:
        print('  aviso: ficheros ajenos en la salida: %s' % sobran)

    idem = {'ejecutada': False}
    if args.idempotencia:
        print('== 2/3 · idempotencia: 2.ª construcción en un clon', flush=True)
        clon = carpeta_dryrun(None) + '-idem'
        if os.path.isdir(clon):
            shutil.rmtree(clon)
        construir(clon, datos, mes, log=lambda *_: None)
        difs = []
        for n in mapas.FICHEROS.values():
            difs += diferencias(digest(os.path.join(carpeta, n)), digest(os.path.join(clon, n)), n)
        shutil.rmtree(clon)
        idem = {'ejecutada': True, 'diferencias': len(difs), 'detalle': difs[:30]}
        print('  diferencias 1.ª vs 2.ª construcción: %d' % len(difs), flush=True)
        for d in difs[:10]:
            print('    ' + d)

    print('== 3/3 · inject_cache (AL FINAL, sobre los 9)', flush=True)
    cache = [] if args.sin_cache else inject_cache(carpeta)
    ver = verificar_cache(carpeta) if cache else {}
    for f, (n, v) in ver.items():
        print('    data_only %-44s %4d fórmulas · %4d con valor' % (f, n, v))

    fallos = []
    if idem.get('diferencias'):
        fallos.append('idempotencia: %d diferencias' % idem['diferencias'])
    fallos += ['inject_cache %s: exit %s' % (c['fichero'], c['exit']) for c in cache if c['exit']]
    fallos += ['inject_cache %s: %s' % (c['fichero'], c['salida']) for c in cache
               if 'fallos_pycel=0' not in c['salida']]
    salida = OrderedDict([('producto', mapas.SLUG), ('version', VERSION),
                          ('fecha', datetime.datetime.now().isoformat(timespec='seconds')),
                          ('modo', 'dry-run' if args.dry_run else 'real'), ('carpeta', carpeta), ('mes', mes),
                          ('libros', informe), ('idempotencia', idem), ('inject_cache', cache),
                          ('cache_data_only', ver), ('fallos', fallos),
                          ('segundos', round(time.time() - t0, 1))])
    if args.json:
        with open(args.json, 'w', encoding='utf-8') as fh:
            json.dump(salida, fh, ensure_ascii=False, indent=1, default=str)
        print('informe → %s' % args.json)
    if fallos:
        print('\nFALLOS:\n  ' + '\n  '.join(fallos))
        sys.exit(1)
    print('\nOK: 9 xlsx en %s (%.0f s)' % (carpeta, time.time() - t0))


if __name__ == '__main__':
    main()
