#!/usr/bin/env python3
"""
semilla_pol.py — Plan de Negocio: Pollería y Asador de Pollos (ES 1.0; F1-camino §1-§3): prepara el ORIGEN que
`planes-v2_0/main.py` convierte en el plan financiero y el checklist de la pollería. Copia adaptada de
`plan-taqueria/semilla_taq.py`.

Copia los dos xlsx de la Taquería ES 1.0 publicados (`dl/plan-negocio-taqueria/`, SOLO plan y checklist, SOLO
LECTURA; molde A-β + checklist C2) con los nombres de fichero de la pollería y escribe lo que el módulo de contenido no
alcanza (F1-camino §1, medido el 29-oct):
  · plan: valores de `0. Supuestos` (grupo_a conserva lo verde que trae el libro; `B60`/`B61` ya son 7/12 en la
    Taquería: assert, sin escribir), títulos de hoja (Inversión A1/A2/A3, PyG A2, Punto Equilibrio A2), las 18
    partidas de inversión (filas 7-24: rótulo, importe, nota, ¿IVA?), los 9 costes fijos propios del P&L (filas
    26-34: rótulo e importe; la nota la pone el contenido) y los docProps description/keywords. `PyG!G13` ya es VALOR
    en la Taquería y lo rehace el motor con `IVA_COMPRAS` (0,091): aquí NO se toca;
  · `_sin_rotacion`: la Taquería trae ENCENDIDO el bloque de rotación y la pollería lo apaga (`ROTACION = {'activa':
    False}`): sin quitar del origen sus 16 fórmulas (`0. Supuestos!B50:C50` y `Punto Equilibrio!B26:D26, A28,
    A37:E38`), `main.py` aborta con «el post-proceso PERDIÓ 15 fórmulas» (F1-camino §1, vuelta 2);
  · checklist: título, pestañas, cabeceras de fase y las tareas de `datos_checklist.py` con el pie «Tareas
    completadas:» debajo; `Instrucciones!A3` con el nombre del producto; DV, formato condicional y contadores los rehace
    el motor.
Corre en el VPS (openpyxl). Después: `main.py --producto plan-negocio-polleria --dry-run --origen <salida>` (ciclo.sh).

    /root/venv-guias/bin/python semilla_pol.py --salida /root/tmp-pol-f2/motor/origen

Sesión Claude Code, 29-oct-2026.
"""
import argparse
import copy
import importlib
import os
import shutil
import sys

import openpyxl

import datos_checklist as DC

AQUI = os.path.dirname(os.path.abspath(__file__))
REPO = os.path.abspath(os.path.join(AQUI, '..', '..', '..'))
TAQ = os.path.join(REPO, 'astro-site', 'public', 'dl', 'plan-negocio-taqueria')
PLAN_TAQ, CHECK_TAQ = 'plan-financiero-taqueria.xlsx', 'checklist-apertura-taqueria.xlsx'
PLAN, CHECK = 'plan-financiero-polleria.xlsx', 'checklist-apertura-polleria.xlsx'
PID = 'plan-negocio-polleria'
NOMBRE = 'Plan de Negocio: Pollería y Asador de Pollos'

TITULOS_PLAN = {
    ('Inversión Inicial', 'A1'): 'PLAN DE NEGOCIO: POLLERÍA Y ASADOR DE POLLOS',
    ('Inversión Inicial', 'A2'): 'Inversión inicial estimada — pollería-asador de barrio con asador de gas, mostrador '
                                 'y 20 plazas',
    ('Inversión Inicial', 'A3'): 'Ajusta las cifras a tu proyecto: son las de un caso tipo en una ciudad media, sin '
                                 'traspaso. Pon tus presupuestos reales en las celdas de importe.',
    ('PyG 3 Años', 'A2'): 'Escenario realista — pollería-asador con mostrador, encargos, plataforma y sala',
    ('Punto Equilibrio', 'A2'): 'Punto de equilibrio mensual y diario — pollería-asador de barrio',
}

#: Filas 7-24 de 'Inversión Inicial': (rótulo, importe, nota, ¿lleva IVA?). R §3: 18 partidas, 88.500 € (obra 46.000 ·
#: maquinaria 33.000 · no amortizable 9.500) + imprevistos del 10 % sobre la obra (fila 27, por fórmula) = 93.100 €.
#: Sin palabras de recuento en el rótulo (motor.RX_RECUENTO), sin paréntesis con número («doce» en letra). ⚠️
#: «campana» a secas la acentúa `motor.TILDES` a «campaña»: siempre «campana extractora»; y nunca el verbo «publica».
INVERSION = [
    ('Proyecto técnico de actividad y de salida de humos, y dirección de obra', 3500,
     'Proyecto de la actividad, del conducto de humos y de la instalación de gas, memoria para la licencia o la '
     'declaración responsable y certificado final de obra', 'Sí'),
    ('Obra y acondicionamiento del local, rótulo y fachada', 24000,
     'Obrador alicatado, suelo antideslizante (también en los 2 m alrededor del asador), aseos (uno adaptado), '
     'mostrador de obra, falsos techos y aislamiento acústico, más el rótulo y la fachada. Estimación del kit: pide '
     'dos presupuestos', 'Sí'),
    ('Conducto de extracción exclusivo hasta cubierta', 6000,
     'Conducto exclusivo hasta la cubierta para el asador, la freidora y el fogón: Madrid pide superar en 1 m lo que '
     'haya en 15 m (Ordenanza 4/2021). Con brasa de leña o carbón, además, filtro de partículas. Ningún fabricante da '
     'precio por metro: pide oferta con tu proyecto', 'Sí'),
    ('Instalación eléctrica trifásica, cuadro y boletín', 4500,
     'Freidora, frío, extracción, climatización, lavavajillas y el motor del asador. Pregunta la potencia disponible '
     'ANTES de firmar: con un asador eléctrico hace falta trifásica, y ampliar la potencia tarda de 4 a 8 semanas',
     'Sí'),
    ('Fontanería, separador de grasas, climatización e instalación de gas', 8000,
     'Agua fría y caliente en el obrador, desagües con separador de grasas, climatización de la sala y la instalación '
     'de gas del asador, la freidora y el fogón por empresa instaladora habilitada, con su certificado (RD 919/2006)',
     'Sí'),
    ('Asador de pollos de gas de doce espadas y juego de espadas', 5000,
     'Asador vertical de gas de doce espadas, de 4 a 6 pollos por espada (60 por carga con 5). En el catálogo de un '
     'distribuidor de hostelería, de 4.373 a 5.967 € sin IVA, con las espadas aparte (2026). Con ocho espadas, de '
     '2.493 a 4.126 €: mira la hoja del asador de la calculadora antes de elegir', 'Sí'),
    ('Vitrina caliente de mostrador y mantenedor de patatas', 2000,
     'Vitrina expositora caliente que mantenga el pollo a 63 °C o más (las genéricas de 85-146 l rondan los '
     '850-1.250 €) y un mantenedor de patatas fritas', 'Sí'),
    ('Freidora doble y fogón', 3000,
     'Freidora de 2 × 10 l para las patatas y las croquetas (unos 1.900 € sin IVA la eléctrica de 15 kW) y un fogón '
     'para las guarniciones y las salsas. Estimación del kit', 'Sí'),
    ('Campana extractora con motor y filtros', 2000,
     'Campana extractora de unos 3 m con motor y filtros de grasa sobre el asador, la freidora y el fogón, con un '
     'caudal de 1.800-2.200 m³/h para doce espadas. Sin precio de referencia para 3 m: estimación del kit', 'Sí'),
    ('Frío: armarios de pollo crudo y adobado, congelador y vitrina refrigerada', 5500,
     'Armarios para el pollo crudo y el adobado a 4 °C o menos (las mollejas, a 3 °C o menos), congelador y la '
     'vitrina refrigerada de las guarniciones del mostrador', 'Sí'),
    ('Lavavajillas de capota, fregaderos, lavamanos y mesas inox de espetar', 5500,
     'Lavavajillas de capota (las espadas caben), fregaderos, lavamanos de accionamiento no manual y una mesa de acero '
     'inoxidable para espetar el pollo crudo, SEPARADA de la de servir', 'Sí'),
    ('Utillaje: sondas, básculas, GN, tijeras de trinchar, etiquetadora y vajilla', 2500,
     'Termómetros sonda, básculas, recipientes GN, tijeras de trinchar, vajilla de la sala y una etiquetadora para lo '
     'que envases antes de que el cliente lo pida (RD 126/2015, art. 5)', 'Sí'),
    ('Botellero y nevera de bebida para llevar', 1500,
     'Botellero de la barra y nevera de bebida a la vista del mostrador. Si la cervecera cede la nevera, esta partida '
     'baja', 'Sí'),
    ('Mobiliario de sala', 4000,
     'Mesas, sillas y taburetes para 16 plazas en mesa y 4 en barra', 'Sí'),
    ('TPV con Veri*factu, datáfono, impresora y tablet de plataformas', 2000,
     'Software de facturación conforme a Veri*factu, datáfono, impresora de tiques y la tablet de las plataformas',
     'Sí'),
    ('Constitución de la SL, licencia, tasas, registro sanitario y APPCC', 3000,
     'SL unipersonal, licencia o declaración responsable, tasas municipales, registro sanitario autonómico y '
     'redacción del plan APPCC. Las tasas no llevan IVA; los honorarios sí: cambia la columna de al lado si tu '
     'presupuesto es sobre todo de honorarios', 'No'),
    ('Marca, web, fotos y lanzamiento', 4000,
     'Identidad, web con la carta, sesión de fotos, ficha de Google Business, tarjetas de encargo y campaña de '
     'apertura', 'Sí'),
    ('Stock inicial: pollo, patata, aceite, envases y bebida', 2500,
     'Pollo, patata, aceite de girasol, especias de los adobos, bolsas térmicas y envases para llevar, y la bebida',
     'Sí'),
]
assert len(INVERSION) == 18 and sum(x[1] for x in INVERSION) == 88500, sum(x[1] for x in INVERSION)

#: Filas 26-34 de 'PyG 3 Años' (costes fijos propios): 9 filas, una por fila del molde (no se suprimen ni se añaden:
#: desplazaría el P&L), 13.900 €/año (R §3: 12.400 €; +1.000 € de mantenimiento —dos limpiezas al año de campana y
#: conducto— y +500 € de marketing = Taquería: F1-camino §1, vuelta 1, cierra la demo §2.11.5). Sin palabras de recuento
#: en el rótulo. Las claves de `contenido….FIJOS` son estos rótulos normalizados.
FIJOS = {
    26: ('Gestoría laboral, fiscal y contable', 2400),
    27: ('Mantenimiento de la campana extractora, el gas, el asador y el frío', 2800),
    28: ('Marketing y publicidad', 3500),
    29: ('Software, TPV, teléfono e internet', 1500),
    30: ('Limpieza, lavandería, EPI y análisis de autocontrol', 1200),
    31: ('Gestión de residuos y del aceite y la grasa usados', 900),
    32: ('Desinsectación, desratización y desinfección (DDD)', 450),
    33: ('Derechos de autor por música ambiental (SGAE y AGEDI-AIE)', 400),
    34: ('Prevención de riesgos laborales, vigilancia de la salud y formación', 750),
}
assert sum(x[1] for x in FIJOS.values()) == 13900

#: Bloque de rotación de la Taquería (F1-camino §1, vuelta 2): 16 celdas con fórmula que el motor ya no escribe con la
#: rotación apagada. `Punto Equilibrio` filas 25-26 (rótulos de la rotación: las reescribe el motor al subir el bloque),
#: A28 (frase de interpretación de la rotación) y A37:E38 (las dos últimas filas de la tabla de sensibilidad, que sube 2).
SIN_ROTACION = (('0. Supuestos', ('B50', 'C50')),
                ('Punto Equilibrio', ('B26', 'C26', 'D26', 'A28') + tuple('%s%d' % (c, r) for r in (37, 38)
                                                                         for c in 'ABCDE')))
assert sum(len(c) for _h, c in SIN_ROTACION) == 16


def _contenido():
    sys.path.insert(0, os.path.join(os.path.dirname(AQUI), 'planes-v2_0'))
    import grupo_a
    return grupo_a, importlib.import_module('contenido_plan_negocio_polleria.a')


def _supuestos(wb):
    """Los valores de `0. Supuestos` del caso pollería: `grupo_a` CONSERVA lo que el libro ya trae en las celdas
    verdes y sólo usa el valor del contenido en una celda vacía; sin este paso el caso seguiría siendo el de la
    taquería (77 × 15,50 € × 300, 55.000 + 118.000 €)."""
    grupo_a, contenido = _contenido()
    coords = dict((t[0], t[1]) for t in grupo_a.SUPUESTOS_BASE)
    ws = wb['0. Supuestos']
    escritas = 0
    for clave, t in contenido.SUPUESTOS.items():
        if t[2] is None:
            continue
        assert clave in coords, clave
        ws[coords[clave]] = t[2]
        escritas += 1
    # F1-camino §2 y §7 F2 (2): las pagas con el Mes 1 en enero ya vienen de la Taquería (7/12)
    assert coords['mes_paga_1'] == 'B60' and coords['mes_paga_2'] == 'B61', (coords['mes_paga_1'], coords['mes_paga_2'])
    assert (ws['B60'].value, ws['B61'].value) == (7, 12), (ws['B60'].value, ws['B61'].value)
    return escritas


def _sin_rotacion(wb):
    """Quita del origen las 16 fórmulas del bloque de rotación de la Taquería (F1-camino §1, vuelta 2)."""
    quitadas = 0
    for hoja, celdas in SIN_ROTACION:
        ws = wb[hoja]
        for c in celdas:
            v = ws[c].value
            assert isinstance(v, str) and v.startswith('='), (hoja, c, v)
            ws[c].value = None
            quitadas += 1
    pe = wb['Punto Equilibrio']
    assert str(pe['A25'].value).lower().startswith('rotaciones al d'), pe['A25'].value
    return quitadas


def _docprops(wb):
    """El motor no reescribe `description` ni `keywords` y heredaba los de la taquería."""
    wb.properties.description = 'aichef.pro/' + PID
    wb.properties.keywords = 'plan de negocio pollería, asador de pollos, pollo a la brasa, AI Chef Pro'


def plan(salida):
    destino = os.path.join(salida, PLAN)
    shutil.copy2(os.path.join(TAQ, PLAN_TAQ), destino)
    wb = openpyxl.load_workbook(destino)
    print('  supuestos escritos:', _supuestos(wb))
    print('  fórmulas de rotación quitadas:', _sin_rotacion(wb))
    _docprops(wb)
    for (hoja, coord), texto in TITULOS_PLAN.items():
        wb[hoja][coord] = texto
    ws = wb['Inversión Inicial']
    for i, (rot, imp, nota, iva) in enumerate(INVERSION):
        r = 7 + i
        assert not str(ws.cell(r, 1).value).lower().startswith(('fianza', 'alquiler', 'imprevistos')), r
        ws.cell(r, 1).value, ws.cell(r, 2).value, ws.cell(r, 4).value, ws.cell(r, 5).value = rot, imp, nota, iva
    assert str(ws['A25'].value).lower().startswith('fianza'), ws['A25'].value
    assert str(ws['A26'].value).lower().startswith('alquiler'), ws['A26'].value
    assert str(ws['A27'].value).lower().startswith('imprevistos'), ws['A27'].value
    ws = wb['PyG 3 Años']
    assert str(ws['A25'].value).startswith('Seguros') and str(ws['A35'].value).startswith('Amortización'), (
        ws['A25'].value, ws['A35'].value)
    for r, (rot, imp) in FIJOS.items():
        ws.cell(r, 1).value, ws.cell(r, 2).value = rot, imp
        ws.cell(r, 6).value = None          # la nota la pone el contenido (FIJOS)
    # `PyG!G13` (IVA soportado de la compra de comida): en la Taquería ya es VALOR (0,085) y el motor lo reescribe con
    # `IVA_COMPRAS` (0,091). No se toca aquí (F1-camino §3).
    assert str(ws['A13'].value).startswith('Coste de mercancía — comida'), ws['A13'].value
    assert not str(ws['G13'].value).startswith('='), ws['G13'].value
    wb.save(destino)
    return destino


def _copiar_estilo(origen, destino):
    destino.font, destino.border = copy.copy(origen.font), copy.copy(origen.border)
    destino.fill, destino.alignment = copy.copy(origen.fill), copy.copy(origen.alignment)
    destino.number_format, destino.protection = origen.number_format, copy.copy(origen.protection)


def checklist(salida):
    destino = os.path.join(salida, CHECK)
    shutil.copy2(os.path.join(TAQ, CHECK_TAQ), destino)
    wb = openpyxl.load_workbook(destino)
    fases = [ws for ws in wb.worksheets if ws.title[:2] in ('F1', 'F2', 'F3', 'F4', 'F5', 'F6')]
    assert len(fases) == 6
    for ws, (titulo, cabecera, tareas) in zip(fases, DC.FASES):
        ws.title = titulo
        ws['A1'], ws['A2'] = DC.TITULO, cabecera
        assert ws['B4'].value and 'TRÁMITE' in ws['B4'].value.upper(), ws['B4'].value
        modelo = [copy.copy(ws.cell(5, c)) for c in range(1, 6)]
        alto = ws.row_dimensions[5].height
        pie = [copy.copy(ws.cell(r, c)) for r in range(5, ws.max_row + 1) for c in range(1, 6)
               if isinstance(ws.cell(r, c).value, str) and 'completad' in ws.cell(r, c).value.lower()]
        assert pie, ws.title
        fila_pie = pie[0].row
        estilo_pie = [copy.copy(ws.cell(fila_pie, c)) for c in range(1, 6)]
        # vaciar la zona de tareas y el pie (valores, DV y formato condicional los rehace el motor)
        for r in range(5, ws.max_row + 1):
            for c in range(1, 6):
                ws.cell(r, c).value = None
            ws.row_dimensions[r].height = None
        ws.data_validations.dataValidation = []
        ws.conditional_formatting = type(ws.conditional_formatting)()
        for i, (tarea, resp, plazo, nota) in enumerate(tareas):
            r = 5 + i
            for c, v in enumerate(('☐', tarea, resp, plazo, nota), start=1):
                cel = ws.cell(r, c)
                _copiar_estilo(modelo[c - 1], cel)
                cel.value = v
            ws.row_dimensions[r].height = alto
        r = 5 + len(tareas) + 1
        for c in range(1, 6):
            _copiar_estilo(estilo_pie[c - 1], ws.cell(r, c))
        ws.cell(r, 2).value = 'Tareas completadas:'
        ws.cell(r, 3).value = '=COUNTIF(A5:A%d,"✓")' % (4 + len(tareas))
        ws.cell(r, 4).value = 'de'
        ws.cell(r, 5).value = '=COUNTIF(B5:B%d,"?*")-COUNTIF(A5:A%d,"N/A")' % (4 + len(tareas), 4 + len(tareas))
    # `motor.asegurar_instrucciones` no reescribe la hoja si ya existe: la de la Taquería nombraba su producto
    ins = wb['Instrucciones']
    assert 'Taquería' in str(ins['A3'].value), ins['A3'].value
    ins['A3'] = 'Este fichero forma parte del «%s» de AI Chef Pro.' % NOMBRE
    _docprops(wb)
    wb.save(destino)
    return destino


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--salida', required=True)
    a = ap.parse_args()
    if os.path.isdir(a.salida):
        shutil.rmtree(a.salida)
    os.makedirs(a.salida)
    print('plan      →', plan(a.salida))
    print('checklist →', checklist(a.salida), '(%d tareas: %s)' % (
        sum(len(t) for _, _, t in DC.FASES), '/'.join(str(len(t)) for _, _, t in DC.FASES)))


if __name__ == '__main__':
    main()
