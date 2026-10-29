#!/usr/bin/env python3
"""
calculadora_pol_es.py — Plan de Negocio: Pollería y Asador de Pollos ES 1.0: `calculadora-pollo-asador-escandallo-polleria.xlsx`.
Hojas `Pollo` · `Asador` · `Escandallo` · `Instrucciones`. Gemelo de código: `plan-taqueria/calculadora_taq_es.py`
(helpers, inject_cache, autotest con pycel y un `calcular()` en Python, protección sin contraseña, verdes desbloqueadas, A4,
docProps AI Chef Pro). Todos los valores por defecto son «estimación del kit, editable» (F1-research-polleria.md §2 y §8).
  · Pollo: peso en canal, precio, sazón y rendimiento del asado → peso asado, merma, coste por pollo y por kg asado, gramos y coste
    de ½ y ¼, pollos para N personas, kilos en canal a comprar y coste de los sobrantes que no se reaprovechan.
  · Asador: espadas × pollos por espada = pollos por tanda; minutos de asado (entrada) + carga y descarga = ciclo; tandas entre el
    primer encendido y la hora límite; capacidad frente a la punta del domingo (y de agosto), encendido más tardío y gas por pollo.
  · Escandallo: 7 guarniciones y salsas + pan + bebida + el pollo (en fracciones) × 9 columnas (entero, ½, ¼ y 6 combos) → coste,
    food cost con semáforo, margen, PVP sugerido y margen por canal (sala · mostrador y encargo · plataforma) y puesto por margen.
Valores que otros scripts leen de la CACHÉ del xlsx (dict `EXPUESTOS`, función `valores_expuestos(ruta)`; celdas por defecto):
    POLLOS_POR_TANDA    Asador!B%(POLLOS_POR_TANDA)s     60 pollos (12 espadas × 5)
    CAPACIDAD_DOMINGO   Asador!B%(CAPACIDAD_DOMINGO)s     180 pollos antes de la hora límite (3 tandas × 60)
    RENDIMIENTO_POLLO   Pollo!B%(RENDIMIENTO_POLLO)s      0,76 (76 %)
    COSTE_POLLO_ENTERO  Pollo!B%(COSTE_POLLO_ENTERO)s      4,5525 € (1,45 kg × 3,05 €/kg + 0,13 € de sazón)
    FC_POLLO_ENTERO     Escandallo!D%(FC_POLLO_ENTERO)s    ≈ 38,8 % (sobre el precio sin IVA)
Corre en el VPS (openpyxl + pycel):

    /root/venv-guias/bin/python calculadora_pol_es.py --salida <dir>      # genera + inject_cache + autotest
    /root/venv-guias/bin/python calculadora_pol_es.py --salida <dir> --autotest

Sesión Claude Code, 29-oct-2026.
"""
import argparse
import math
import os
import re
import sys
from collections import OrderedDict

import openpyxl
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill, Protection
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

AQUI = os.path.dirname(os.path.abspath(__file__))
PD = os.path.dirname(AQUI)
sys.path.insert(0, AQUI)
import pol_version                                                                 # noqa: E402,F401  (parchea motor: 1.0 · noviembre 2026)
import motor                                                                       # noqa: E402

PID = 'plan-negocio-polleria'
FICHERO = 'calculadora-pollo-asador-escandallo-polleria.xlsx'
TITULO = 'Calculadora del pollo, el asador y el escandallo'
HP, HA, HE, HI = 'Pollo', 'Asador', 'Escandallo', 'Instrucciones'
VERDE = PatternFill('solid', fgColor='E8F5E9')
GRIS = PatternFill('solid', fgColor='F2F2F2')
ROJO = PatternFill('solid', bgColor='F8CBAD', fgColor='F8CBAD')
AMBAR = PatternFill('solid', bgColor='FFE699', fgColor='FFE699')
OK = PatternFill('solid', bgColor='C6E0B4', fgColor='C6E0B4')
FUENTE = 'Calibri'
EUR, ENT, PCT, DEC, GR, EUR3 = '#,##0.00 €', '#,##0', '0.0%', '#,##0.0', '#,##0.0', '#,##0.000 €'
FC_OK, FC_VIG, FC_ALTO = 'En objetivo', 'Vigila el coste', 'Por encima del objetivo'
SEM_OK, SEM_KO = 'Cubierto', 'No llegas a la punta'
EST = ' Estimación del kit, editable.'

# ---------------------------------------------------------------------------------------------------------------------
# Datos (F1-research §2 y §8). Importes sin IVA salvo el PVP de carta.
# ---------------------------------------------------------------------------------------------------------------------
#: hoja Pollo: (clave, rótulo, formato, valor por defecto, constante gris, nota)
ENT_P = [
    ('P', 'Peso medio del pollo en canal (kg)', '#,##0.00', 1.45, False, 'Pollo limpio de 1,4 a 1,5 kg: el calibre que pide el asador. Pésalo en la recepción.' + EST),
    ('Pk', 'Precio del pollo en canal (€/kg, sin IVA)', EUR, 3.05, False, 'Unos 2,96 € el kilo en origen y 3,25 € en Makro (octubre de 2026): el kit pone 3,05 €. Pon el de tu proveedor.' + EST),
    ('S', 'Sazón, aceite y sal por pollo (€, sin IVA)', EUR, 0.13, False, 'Adobo, aceite y sal de un pollo.' + EST),
    ('r', 'Rendimiento del asado (peso asado ÷ peso en canal)', PCT, 0.76, False,
     'Lo que queda del pollo tras el asado, con hueso. Mídelo con 5 pollos de la misma carga (hoja «Instrucciones»).' + EST),
    ('ref78', 'Referencia: rendimiento del pollo entero asado (USDA)', '0%', 0.78, True,
     'Referencia fija: el USDA da 78 % con un rango de 68 a 84 %. No se edita.'),
    ('pp', 'Personas que comen de un pollo', ENT, 3, False, 'Número entero mayor que 0. Para un plato principal con guarnición suelen ser 3.' + EST),
    ('Np', 'Personas que quieres atender (encargo o evento)', ENT, 40, False, 'Número entero mayor que 0. Te dice cuántos pollos pedir.'),
    ('Ns', 'Pollos sobrantes al cierre', DEC, 2.0, False, 'Pollos enteros que quedan sin vender al cerrar. Cuéntalos; por defecto 2.' + EST),
    ('a', 'Parte del sobrante que se reaprovecha (%)', PCT, 0.50, False,
     'Croquetas, ensalada de pollo o relleno del día siguiente, tras el enfriado correcto. La parte que se tira cuesta dinero.' + EST),
    ('KG', 'Gramos por kilo (constante)', ENT, 1000, True, 'Constante: un kilo son 1.000 gramos. No se edita.'),
    ('M2', 'Medios pollos que salen de un pollo (constante)', ENT, 2, True, 'Constante: un pollo se parte en 2 mitades. No se edita.'),
    ('M4', 'Cuartos de pollo que salen de un pollo (constante)', ENT, 4, True, 'Constante: un pollo se parte en 4 cuartos. No se edita.'),
]
SAL_P = [
    ('asado', 'PESO ASADO DE UN POLLO (kg)', '#,##0.000', '{P}*{r}', 'Peso en canal × rendimiento.'),
    ('merma_kg', 'Merma del asado (kg)', '#,##0.000', '{P}-{asado}', 'Agua, grasa y piel que se pierden en el asador.'),
    ('merma_pct', 'Merma del asado (%)', PCT, '{merma_kg}/{P}', 'Merma ÷ peso en canal.'),
    ('rend', 'RENDIMIENTO DEL POLLO (%)', PCT, '{asado}/{P}', 'Peso asado ÷ peso en canal.'),
    ('vs_ref', 'Tu rendimiento frente a la referencia USDA (%)', '0%', '{rend}/{ref78}', 'Por debajo de 100 %: pierdes más que la referencia; mídelo otra vez.'),
    ('coste', 'COSTE POR POLLO ENTERO (€)', EUR3, '{P}*{Pk}+{S}', 'Peso en canal × precio por kilo + sazón. Alimenta el escandallo.'),
    ('coste_kg', 'COSTE POR KG ASADO (€)', EUR, '{coste}/{asado}', 'Coste del pollo ÷ kilos asados.'),
    ('g_entero', 'Gramos asados del pollo entero (g)', GR, '{asado}*{KG}', 'Peso asado × 1.000.'),
    ('g_medio', 'Gramos asados de medio pollo (g)', GR, '{g_entero}/{M2}', 'Pollo entero asado ÷ 2.'),
    ('g_cuarto', 'Gramos asados de un cuarto (g)', GR, '{g_entero}/{M4}', 'Pollo entero asado ÷ 4.'),
    ('c_medio', 'Coste de medio pollo (€)', EUR3, '{coste}/{M2}', 'Coste del pollo ÷ 2.'),
    ('c_cuarto', 'Coste de un cuarto de pollo (€)', EUR3, '{coste}/{M4}', 'Coste del pollo ÷ 4.'),
    ('pollos_N', 'POLLOS PARA LAS N PERSONAS (pollos)', ENT, 'ROUNDUP({Np}/{pp},0)', 'Personas ÷ personas por pollo, redondeado hacia arriba.'),
    ('kg_canal', 'Kilos en canal que tienes que comprar (kg)', '#,##0.00', '{pollos_N}*{P}', 'Pollos × peso en canal.'),
    ('c_merma', 'Coste de los sobrantes que se tiran al cierre (€)', EUR, '{Ns}*(1-{a})*{coste}', 'Sobrantes × parte que no se reaprovecha × coste del pollo.'),
]

#: hoja Asador
ENT_S = [
    ('E', 'Espadas del asador', ENT, 12, False, 'Número entero mayor que 0. Pon las de tu modelo (los hay de 6, 8 y 12).'),
    ('p', 'Pollos por espada', ENT, 5, False, 'Las fichas dan de 4 a 6 según el calibre. Número entero mayor que 0.'),
    ('t', 'Minutos de asado de una tanda', DEC, 75.0, False,
     'SIN FUENTE PUBLICADA: ninguna ficha da el tiempo de un pollo de 1,5 kg. Mídelo con tu asador (hoja «Instrucciones»).' + EST),
    ('c', 'Minutos de carga, descarga y limpieza entre tandas', DEC, 15.0, False, 'Espetar la tanda siguiente mientras sale la anterior.' + EST),
    ('hh0', 'Primer encendido: hora', ENT, 10, False, 'Hora entera (0 a 23). Con 10 y 30 minutos, el asador se enciende a las 10:30.' + EST),
    ('mm0', 'Primer encendido: minutos', ENT, 30, False, 'Minutos (0 a 59).'),
    ('hh1', 'Hora límite para tener los pollos asados: hora', ENT, 15, False, 'La punta del domingo termina hacia las 15:00. Hora entera (0 a 23).' + EST),
    ('mm1', 'Hora límite para tener los pollos asados: minutos', ENT, 0, False, 'Minutos (0 a 59).'),
    ('D', 'Pollos que vendes un domingo normal', ENT, 106, False, 'El del caso del plan de negocio. Cámbialo por el tuyo.' + EST),
    ('hp', 'Parte del domingo que se vende en la punta (%)', PCT, 0.80, False, 'Entre las 12:30 y las 15:00.' + EST),
    ('ag', 'Subida del domingo de agosto o de un festivo (%)', PCT, 0.15, False, 'Sobre un domingo normal. En la costa puede ser mayor.' + EST),
    ('kw', 'Potencia media de gas que consume el asador (kW)', DEC, 60.0, False,
     'La ficha de un asador de 8 espadas da 56 kW a pleno rendimiento; el de 12 espadas gasta más y su media real queda por debajo del máximo de su ficha. Pon la de la tuya.' + EST),
    ('ekwh', 'Precio del gas (€/kWh, sin IVA)', '#,##0.000 €', 0.10, False, 'Butano o propano. Pon el de tu contrato o tu bombona.' + EST),
    ('MH', 'Minutos por hora (constante)', ENT, 60, True, 'Constante: una hora son 60 minutos. No se edita.'),
]
SAL_S = [
    ('carga', 'POLLOS POR TANDA', ENT, '{E}*{p}', 'Espadas × pollos por espada.'),
    ('ciclo', 'Minutos de un ciclo completo', DEC, '{t}+{c}', 'Asado + carga y descarga.'),
    ('tandas_h', 'Tandas por hora', '#,##0.00', '{MH}/{ciclo}', '60 minutos ÷ minutos del ciclo.'),
    ('cap_h', 'Capacidad por hora (pollos)', DEC, '{carga}*{tandas_h}', 'Pollos por tanda × tandas por hora.'),
    ('min_disp', 'Minutos entre el primer encendido y la hora límite', ENT, '({hh1}-{hh0})*{MH}+({mm1}-{mm0})', 'Si sale negativo, la hora límite es anterior al encendido.'),
    ('tandas', 'TANDAS ANTES DE LA HORA LÍMITE', ENT, 'INT(MAX(0,{min_disp})/{ciclo})', 'Ciclos completos que caben en ese tiempo.'),
    ('cap_dom', 'CAPACIDAD DEL DOMINGO ANTES DE LA HORA LÍMITE (pollos)', ENT, '{carga}*{tandas}', 'Pollos por tanda × tandas.'),
    ('dem_punta', 'Pollos que se venden en la punta de un domingo normal', DEC, '{D}*{hp}', 'Pollos del domingo × parte que se vende en la punta.'),
    ('cargas_dom', 'Tandas que necesita un domingo normal (en todo el día)', ENT, 'ROUNDUP({D}/{carga},0)', 'Pollos del domingo ÷ pollos por tanda, hacia arriba.'),
    ('cargas_punta', 'Tandas que necesita la punta de un domingo normal', ENT, 'ROUNDUP({dem_punta}/{carga},0)', 'Pollos de la punta ÷ pollos por tanda, hacia arriba.'),
    ('balance', 'Sobrante (+) o déficit (−) en la punta (pollos)', DEC, '{cap_dom}-{dem_punta}', 'Capacidad antes de la hora límite menos pollos de la punta.'),
    ('sem', 'Semáforo de la punta (domingo normal)', 'General', 'IF({balance}>=0,"%s","%s")' % (SEM_OK, SEM_KO), 'Cubierto si la capacidad llega a los pollos de la punta.'),
    ('d_ago', 'Pollos de un domingo de agosto o festivo', ENT, 'ROUNDUP({D}*(1+{ag}),0)', 'Domingo normal más la subida, hacia arriba.'),
    ('cargas_ago', 'Tandas que necesita un domingo de agosto o festivo', ENT, 'ROUNDUP({d_ago}/{carga},0)', 'Pollos de agosto ÷ pollos por tanda, hacia arriba.'),
    ('balance_ago', 'Sobrante (+) o déficit (−) en la punta de agosto (pollos)', DEC, '{cap_dom}-{d_ago}*{hp}', 'Capacidad menos la punta de agosto.'),
    ('sem_ago', 'Semáforo de la punta de agosto', 'General', 'IF({balance_ago}>=0,"%s","%s")' % (SEM_OK, SEM_KO), 'Cubierto si la capacidad llega a la punta de agosto.'),
    ('tarde_min', 'Encendido más tardío (minutos desde las 00:00)', ENT, '{hh1}*{MH}+{mm1}-{cargas_punta}*{ciclo}', 'Hora límite menos las tandas de la punta por su ciclo.'),
    ('tarde_h', 'Encendido más tardío: hora', ENT, 'INT({tarde_min}/{MH})', 'Si enciendes después, no llegas a la punta.'),
    ('tarde_m', 'Encendido más tardío: minutos', ENT, 'MOD({tarde_min},{MH})', ''),
    ('horas_gas', 'Horas de asador encendido en un domingo normal', '#,##0.00', '{cargas_dom}*{ciclo}/{MH}', 'Tandas × ciclo ÷ 60.'),
    ('gas_dia', 'Gas de un domingo normal (€)', EUR, '{kw}*{horas_gas}*{ekwh}', 'Potencia × horas × precio por kWh.'),
    ('gas_pollo', 'GAS POR POLLO ASADO (€)', EUR3, '{gas_dia}/({cargas_dom}*{carga})', 'Gas del día ÷ pollos que caben en las tandas del día.'),
]

#: hoja Escandallo
ESC_IN = [
    ('iva', 'IVA de la comida (%)', PCT, 0.10, 'Comida preparada, en sala, para llevar y a domicilio: 10 % (Ley 37/1992, art. 91).'),
    ('ivab', 'IVA de la bebida que sale del local (%)', PCT, 0.21, 'Cerveza, vino y refrescos con azúcares o edulcorantes añadidos para llevar: 21 % (Ley 37/1992, art. 91).'),
    ('ivap', 'IVA del pan vendido suelto (%)', PCT, 0.04, 'Pan común: 4 % (Ley 37/1992, art. 91).'),
    ('obj', 'Food cost objetivo (% sobre el precio sin IVA)', PCT, 0.33, 'Referencia del pedido tipo para llevar: 33 %.' + EST),
    ('tol', 'Margen de vigilancia sobre el objetivo (puntos)', PCT, 0.06, 'Hasta este margen por encima del objetivo, el semáforo dice «Vigila el coste». El pollo entero suele caer aquí.'),
    ('merm', 'Merma de producción de las guarniciones (%)', PCT, 0.03, 'Patata frita que se tira, ensalada que no se aprovecha. El pollo no la lleva: su merma está en el rendimiento.' + EST),
    ('env', 'Envase por pedido para llevar (€, sin IVA)', EUR, 0.60, 'Bolsa térmica, bandeja y tarrinas de salsa de un pedido. El envase de plástico se cobra aparte en el tique.' + EST),
    ('com', 'Comisión de la plataforma de reparto (%)', PCT, 0.30, 'Sobre el precio sin IVA del pedido. Revisa tu contrato.' + EST),
    ('rec', 'Recargo de precio en la plataforma (%)', PCT, 0.15, 'Lo que subes el precio en la plataforma para compensar (unos 2 € por pieza).' + EST),
]
#: gramos: (ingrediente, precio €/kg, origen del precio)
PF, PA, CRO, ENS, ENL, ALI, YOG = ('Patatas fritas (congeladas, con aceite)', 'Patatas asadas (frescas)', 'Croquetas de pollo (con el sobrante)',
                                   'Ensaladilla (con ovoproducto)', 'Ensalada mixta', 'Alioli (con ovoproducto pasteurizado)',
                                   'Salsa de yogur y hierbas (sin huevo)')
ING_G = [
    (PF, 2.55, 'estimación del kit'),
    (PA, 1.10, 'estimación del kit'),
    (CRO, 8.50, 'estimación del kit (croquetas de 25 g hechas con el pollo del día anterior)'),
    (ENS, 6.25, 'estimación del kit'),
    (ENL, 4.00, 'estimación del kit'),
    (ALI, 7.00, 'estimación del kit'),
    (YOG, 4.50, 'estimación del kit'),
]
#: unidades: (ingrediente, precio €/ud o None si es el coste del pollo de la hoja Pollo, origen)
UPOL, UPAN, UBEB = 'Pollo asado (fracción de pollo entero)', 'Pan (ud)', 'Bebida (lata o botellín, ud)'
ING_U = [
    (UPOL, None, 'enlazado desde la hoja Pollo (coste por pollo entero)'),
    (UPAN, 0.18, 'estimación del kit'),
    (UBEB, 0.55, 'estimación del kit'),
]
#: columna → (PVP con IVA, de ello bebida que sale del local con IVA, de ello pan suelto con IVA, {ingrediente en g: g}, {unidades: cantidad})
RECETAS = OrderedDict([
    ('Pollo entero', (12.90, 0.0, 0.0, {}, {UPOL: 1})),
    ('Medio pollo', (7.00, 0.0, 0.0, {}, {UPOL: 0.5})),
    ('Cuarto de pollo', (4.00, 0.0, 0.0, {}, {UPOL: 0.25})),
    ('Pollo con patatas', (16.40, 0.0, 0.0, {PF: 350, ALI: 25}, {UPOL: 1})),
    ('Medio pollo con croquetas y ensalada', (12.50, 0.0, 0.0, {CRO: 150, ENL: 100}, {UPOL: 0.5})),
    ('Menú de cuarto de sala', (9.50, 0.0, 0.0, {PF: 250, ENL: 70}, {UPOL: 0.25, UPAN: 1})),
    ('Medio a la brasa con papas, ensalada y cremas', (14.50, 0.0, 0.0, {PA: 300, ENL: 120, ALI: 25, YOG: 30}, {UPOL: 0.5})),
    ('Familiar de domingo', (29.90, 0.0, 0.0, {PF: 500, CRO: 160, ENS: 200, ENL: 150, ALI: 40, YOG: 40}, {UPOL: 1, UPAN: 2})),
    ('Combo de plataforma', (21.90, 1.80, 0.0, {PA: 400, ENL: 100, YOG: 40}, {UPOL: 1, UPAN: 1, UBEB: 1})),
])
NP, NG, NU = len(RECETAS), len(ING_G), len(ING_U)
PCOL = [get_column_letter(4 + j) for j in range(NP)]                              # D..L
ORI = get_column_letter(4 + NP)                                                    # M: origen del precio
J_ENTERO = 0

# ---------------------------------------------------------------------------------------------------------------------
# Disposición de las hojas (filas)
# ---------------------------------------------------------------------------------------------------------------------
def _filas(ent, sal, h0=5):
    m = OrderedDict((k[0], h0 + 1 + i) for i, k in enumerate(ent))
    ra = h0 + 1 + len(ent) + 1
    s = OrderedDict((k[0], ra + 1 + i) for i, k in enumerate(sal))
    return m, ra, s


M_P, M_RP, S_P = _filas(ENT_P, SAL_P)                                              # hoja Pollo
M_S, M_RS, S_S = _filas(ENT_S, SAL_S)                                              # hoja Asador
E = OrderedDict((k[0], 5 + i) for i, k in enumerate(ESC_IN))
E_CAB = 5 + len(ESC_IN) + 1
E_PVP, E_BEB, E_PAN = E_CAB + 1, E_CAB + 2, E_CAB + 3
E_G0 = E_PAN + 1
E_G1 = E_G0 + NG - 1
E_UH = E_G1 + 1
E_U0 = E_UH + 1
E_U1 = E_U0 + NU - 1
E_R = E_U1 + 2
RES = [
    ('peso', 'Peso servido en el plato (g)', GR),
    ('c_pollo', 'Coste del pollo (€)', EUR3),
    ('c_acomp', 'Coste de guarniciones, pan y bebida, sin merma (€)', EUR3),
    ('c_tot', 'COSTE CON MERMA DE GUARNICIONES (€)', EUR3),
    ('pvp_sin', 'Precio de venta sin IVA (€)', EUR),
    ('fc', 'FOOD COST (%)', PCT),
    ('sem', 'Semáforo contra el objetivo', 'General'),
    ('margen', 'Margen bruto (€, precio sin IVA − coste)', EUR),
    ('pvp_sug', 'PVP con IVA sugerido para llegar al objetivo (€)', EUR),
]
R = OrderedDict((k, E_R + 1 + i) for i, (k, *_r) in enumerate(RES))
E_CH = E_R + len(RES) + 2
CANAL = [
    ('sala_e', 'Sala: margen (€)', EUR), ('sala_p', 'Sala: margen (%)', PCT),
    ('mostr_e', 'Mostrador y encargo: margen (€, descontado el envase)', EUR), ('mostr_p', 'Mostrador y encargo: margen (%)', PCT),
    ('plat_pv', 'Plataforma: precio sin IVA con recargo (€)', EUR),
    ('plat_e', 'Plataforma: margen (€, descontados envase y comisión)', EUR), ('plat_p', 'Plataforma: margen (%)', PCT),
    ('rank', 'Puesto por margen en mostrador (1 = el que más deja)', ENT),
]
CH = OrderedDict((k, E_CH + 1 + i) for i, (k, *_r) in enumerate(CANAL))
E_FIN = E_CH + len(CANAL) + 2

EXPUESTOS = OrderedDict([
    ('POLLOS_POR_TANDA', (HA, 'B%d' % S_S['carga'])),
    ('CAPACIDAD_DOMINGO', (HA, 'B%d' % S_S['cap_dom'])),
    ('RENDIMIENTO_POLLO', (HP, 'B%d' % S_P['rend'])),
    ('COSTE_POLLO_ENTERO', (HP, 'B%d' % S_P['coste'])),
    ('FC_POLLO_ENTERO', (HE, '%s%d' % (PCOL[J_ENTERO], R['fc']))),
])
for _n, (_h, _c) in EXPUESTOS.items():
    __doc__ = __doc__.replace('%(' + _n + ')s', re.sub(r'^[A-Z]+', '', _c))


def valores_defecto():
    """Todas las entradas por defecto en un solo dict plano (claves = las de las celdas verdes)."""
    v = OrderedDict()
    for ent in (ENT_P, ENT_S):
        for k, _r, _f2, val, _c, _n in ent:
            v[k] = val
    for k, _r, _f2, val, _n in ESC_IN:
        v[k] = val
    for i, (_n, p, _o) in enumerate(ING_G):
        v['p%d' % i] = p
    for i, (_n, p, _o) in enumerate(ING_U):
        if p is not None:
            v['up%d' % i] = p
    for j, (_nom, (pvp, beb, pan, gr, ud)) in enumerate(RECETAS.items()):
        v['pvp%d' % j], v['beb%d' % j], v['pan%d' % j] = pvp, beb, pan
        for i, (ni, *_r) in enumerate(ING_G):
            v['g%d_%d' % (i, j)] = float(gr.get(ni, 0))
        for i, (ni, *_r) in enumerate(ING_U):
            v['u%d_%d' % (i, j)] = float(ud.get(ni, 0))
    return v


# ---------------------------------------------------------------------------------------------------------------------
# Cálculo en Python (espejo de las fórmulas) — fuente del autotest
# ---------------------------------------------------------------------------------------------------------------------
def calcular(v):
    r = OrderedDict()
    p = OrderedDict()
    p['asado'] = v['P'] * v['r']
    p['merma_kg'] = v['P'] - p['asado']
    p['merma_pct'] = p['merma_kg'] / v['P']
    p['rend'] = p['asado'] / v['P']
    p['vs_ref'] = p['rend'] / v['ref78']
    p['coste'] = v['P'] * v['Pk'] + v['S']
    p['coste_kg'] = p['coste'] / p['asado']
    p['g_entero'] = p['asado'] * v['KG']
    p['g_medio'] = p['g_entero'] / v['M2']
    p['g_cuarto'] = p['g_entero'] / v['M4']
    p['c_medio'] = p['coste'] / v['M2']
    p['c_cuarto'] = p['coste'] / v['M4']
    p['pollos_N'] = math.ceil(v['Np'] / v['pp'])
    p['kg_canal'] = p['pollos_N'] * v['P']
    p['c_merma'] = v['Ns'] * (1 - v['a']) * p['coste']
    r['P'] = p
    s = OrderedDict()
    s['carga'] = v['E'] * v['p']
    s['ciclo'] = v['t'] + v['c']
    s['tandas_h'] = v['MH'] / s['ciclo']
    s['cap_h'] = s['carga'] * s['tandas_h']
    s['min_disp'] = (v['hh1'] - v['hh0']) * v['MH'] + (v['mm1'] - v['mm0'])
    s['tandas'] = math.floor(max(0, s['min_disp']) / s['ciclo'])
    s['cap_dom'] = s['carga'] * s['tandas']
    s['dem_punta'] = v['D'] * v['hp']
    s['cargas_dom'] = math.ceil(v['D'] / s['carga'])
    s['cargas_punta'] = math.ceil(s['dem_punta'] / s['carga'])
    s['balance'] = s['cap_dom'] - s['dem_punta']
    s['sem'] = SEM_OK if s['balance'] >= 0 else SEM_KO
    s['d_ago'] = math.ceil(v['D'] * (1 + v['ag']))
    s['cargas_ago'] = math.ceil(s['d_ago'] / s['carga'])
    s['balance_ago'] = s['cap_dom'] - s['d_ago'] * v['hp']
    s['sem_ago'] = SEM_OK if s['balance_ago'] >= 0 else SEM_KO
    s['tarde_min'] = v['hh1'] * v['MH'] + v['mm1'] - s['cargas_punta'] * s['ciclo']
    s['tarde_h'] = math.floor(s['tarde_min'] / v['MH'])
    s['tarde_m'] = s['tarde_min'] - v['MH'] * math.floor(s['tarde_min'] / v['MH'])
    s['horas_gas'] = s['cargas_dom'] * s['ciclo'] / v['MH']
    s['gas_dia'] = v['kw'] * s['horas_gas'] * v['ekwh']
    s['gas_pollo'] = s['gas_dia'] / (s['cargas_dom'] * s['carga'])
    r['S'] = s
    r['pg'] = [v['p%d' % i] for i in range(NG)]
    r['pu'] = [p['coste']] + [v['up%d' % i] for i in range(1, NU)]
    cols = []
    for j in range(NP):
        gs = [v['g%d_%d' % (i, j)] for i in range(NG)]
        us = [v['u%d_%d' % (i, j)] for i in range(NU)]
        q = OrderedDict()
        q['peso'] = sum(gs) + us[0] * p['g_entero']
        q['c_pollo'] = us[0] * r['pu'][0]
        q['c_acomp'] = sum(g_ * pr for g_, pr in zip(gs, r['pg'])) / v['KG'] + sum(u_ * pr for u_, pr in zip(us[1:], r['pu'][1:]))
        q['c_tot'] = q['c_pollo'] + q['c_acomp'] * (1 + v['merm'])
        beb, pan = v['beb%d' % j], v['pan%d' % j]
        q['pvp_sin'] = (v['pvp%d' % j] - beb - pan) / (1 + v['iva']) + beb / (1 + v['ivab']) + pan / (1 + v['ivap'])
        q['fc'] = q['c_tot'] / q['pvp_sin']
        q['sem'] = FC_OK if q['fc'] <= v['obj'] else (FC_VIG if q['fc'] <= v['obj'] + v['tol'] else FC_ALTO)
        q['margen'] = q['pvp_sin'] - q['c_tot']
        q['pvp_sug'] = q['c_tot'] / v['obj'] * (1 + v['iva'])
        q['sala_e'] = q['margen']
        q['sala_p'] = q['sala_e'] / q['pvp_sin']
        q['mostr_e'] = q['sala_e'] - v['env']
        q['mostr_p'] = q['mostr_e'] / q['pvp_sin']
        q['plat_pv'] = q['pvp_sin'] * (1 + v['rec'])
        q['plat_e'] = q['plat_pv'] * (1 - v['com']) - q['c_tot'] - v['env']
        q['plat_p'] = q['plat_e'] / q['plat_pv']
        cols.append(q)
    for q in cols:
        q['rank'] = 1 + sum(1 for o in cols if o['mostr_e'] > q['mostr_e'])
    for j, q in enumerate(cols):
        r['p%d' % j] = q
    return r


# ---------------------------------------------------------------------------------------------------------------------
# Helpers (de la Taquería)
# ---------------------------------------------------------------------------------------------------------------------
def _f(ws, coord, v, b=False, sz=11, i=False, fmt=None, wrap=False, fill=None, locked=True):
    c = ws[coord]
    c.value = v
    c.font = Font(name=FUENTE, size=sz, bold=b, italic=i)
    if fmt:
        c.number_format = fmt
    c.alignment = Alignment(wrap_text=wrap, vertical='top')
    if fill:
        c.fill = fill
    c.protection = Protection(locked=locked)
    return c


def _nota(ws, fila, texto, desde='E', hasta='F', chars=80.0):
    if not texto:
        return
    _f(ws, '%s%d' % (desde, fila), texto, wrap=True)
    if desde != hasta:
        ws.merge_cells('%s%d:%s%d' % (desde, fila, hasta, fila))
    ws.row_dimensions[fila].height = max(16, 15 * math.ceil(len(texto) / chars))


def _imprimir(ws, horizontal=False):
    ws.page_setup.paperSize = 9                                                   # A4
    ws.page_setup.orientation = 'landscape' if horizontal else 'portrait'
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.oddFooter.center.text = 'aichef.pro/%s · página &P de &N' % PID
    ws.print_options.horizontalCentered = True
    ws.protection.sheet = True                                                    # sin contraseña
    ws.sheet_view.showGridLines = False


def _dv(ws, tipo, celdas, f1=None, f2=None, op=None, titulo='Valor no válido', msg=''):
    dv = DataValidation(type=tipo, formula1=f1, formula2=f2, operator=op, allow_blank=False, showErrorMessage=True,
                        errorTitle=titulo, error=msg)
    ws.add_data_validation(dv)
    for c in celdas:
        dv.add(c)


def a_excel(expr, ref):
    return re.sub(r'\{(\w+)\}', lambda m: ref[m.group(1)], expr)


def _kg_ext():
    return "'%s'!$B$%d" % (HP, M_P['KG'])


def _bloque(ws, titulo_in, ent, m, ra, sal, s, neg, chars=105.0):
    """Escribe el bloque de entradas (verdes/grises) y el de resultados de una hoja de dos bloques."""
    _f(ws, 'A5', titulo_in, b=True, fill=GRIS)
    _f(ws, 'B5', 'Valor', b=True, fill=GRIS)
    _f(ws, 'C5', 'Notas', b=True, fill=GRIS)
    for k, rot, fmt, val, const, nota in ent:
        f = m[k]
        _f(ws, 'A%d' % f, rot)
        if const:
            _f(ws, 'B%d' % f, val, fmt=fmt, fill=GRIS)
        else:
            _f(ws, 'B%d' % f, val, fmt=fmt, fill=VERDE, locked=False)
        _nota(ws, f, nota, desde='C', hasta='I', chars=chars)
    ref = dict((k, '$B$%d' % f) for k, f in m.items())
    ref.update((k, '$B$%d' % f) for k, f in s.items())
    for col, t in zip('ABC', ('RESULTADOS', 'Valor', 'Notas')):
        _f(ws, '%s%d' % (col, ra), t, b=True, fill=GRIS)
    for k, rot, fmt, expr, nota in sal:
        f = s[k]
        _f(ws, 'A%d' % f, rot, b=k in neg)
        _f(ws, 'B%d' % f, '=' + a_excel(expr, ref), fmt=fmt, b=k in neg, wrap=(fmt == 'General'))
        _nota(ws, f, nota, desde='C', hasta='I', chars=chars)


def _semaforo(ws, celda, ok, ko):
    ws.conditional_formatting.add(celda, FormulaRule(formula=['$%s$%s="%s"' % (celda[0], celda[1:], ok)], fill=OK))
    ws.conditional_formatting.add(celda, FormulaRule(formula=['$%s$%s="%s"' % (celda[0], celda[1:], ko)], fill=ROJO))


# ---------------------------------------------------------------------------------------------------------------------
# Hojas
# ---------------------------------------------------------------------------------------------------------------------
def hoja_pollo(wb):
    ws = wb.active
    ws.title = HP
    for col, w in (('A', 58), ('B', 15), ('C', 15), ('D', 16), ('E', 14), ('F', 14), ('G', 14), ('H', 12), ('I', 15)):
        ws.column_dimensions[col].width = w
    _f(ws, 'A1', 'CALCULADORA DEL POLLO — CUÁNTO PESA ASADO, CUÁNTO CUESTA Y CUÁNTOS POLLOS NECESITAS', b=True, sz=14)
    _f(ws, 'A2', 'Escribe solo en las celdas VERDES: todo lo demás se calcula. Importes sin IVA. Todos los valores por defecto son una '
                 'estimación del kit, editable: cámbialos por los de tu proveedor y, sobre todo, por lo que midas en tus propios pollos.', wrap=True)
    ws.merge_cells('A2:I2')
    ws.row_dimensions[2].height = 32
    _f(ws, 'A3', 'AVISO FIJO. El rendimiento del asado (76 %) es una estimación del kit; la referencia del USDA es 78 %, con un rango de 68 a 84 %. '
                 'Pesa 5 pollos crudos y asados de la misma carga, y ajusta el rendimiento hasta que coincida con lo que mides.',
       wrap=True, b=True, fill=GRIS)
    ws.merge_cells('A3:I3')
    ws.row_dimensions[3].height = 36
    _bloque(ws, 'A. EL POLLO', ENT_P, M_P, M_RP, SAL_P, S_P, ('asado', 'rend', 'coste', 'coste_kg', 'pollos_N'))
    cel = lambda k: 'B%d' % M_P[k]
    _dv(ws, 'decimal', [cel(k) for k in ('P', 'Pk')], f1='0', op='greaterThan', msg='Escribe un número mayor que 0.')
    _dv(ws, 'decimal', [cel(k) for k in ('S', 'Ns')], f1='0', op='greaterThanOrEqual', msg='Escribe un número mayor o igual que 0.')
    _dv(ws, 'decimal', [cel('r')], f1='0.3', f2='1', op='between', titulo='Rendimiento no válido', msg='Escribe entre 0,3 y 1 (0,76 = 76 %).')
    _dv(ws, 'decimal', [cel('a')], f1='0', f2='1', op='between', titulo='Porcentaje no válido', msg='Escribe un porcentaje entre 0 y 1 (0,5 = 50 %).')
    _dv(ws, 'whole', [cel('pp'), cel('Np')], f1='1', f2='1000000', op='between', titulo='Número no válido', msg='Escribe un número entero de 1 o más.')
    _imprimir(ws, horizontal=True)


def hoja_asador(wb):
    ws = wb.create_sheet(HA)
    for col, w in (('A', 62), ('B', 15), ('C', 15), ('D', 16), ('E', 14), ('F', 14), ('G', 14), ('H', 12), ('I', 15)):
        ws.column_dimensions[col].width = w
    _f(ws, 'A1', 'CALCULADORA DEL ASADOR — SI TUS ESPADAS LLEGAN A LA PUNTA DEL DOMINGO', b=True, sz=14)
    _f(ws, 'A2', 'Escribe solo en las celdas VERDES: todo lo demás se calcula. Todos los valores por defecto son una estimación del kit, '
                 'editable. La capacidad cuenta las tandas completas que caben entre el primer encendido y la hora límite.', wrap=True)
    ws.merge_cells('A2:I2')
    ws.row_dimensions[2].height = 32
    _f(ws, 'A3', 'AVISO FIJO. El tiempo de asado (75 minutos) no tiene fuente publicada: ninguna ficha da el de un pollo de 1,5 kg. Cronométralo '
                 'con tu asador y tu calibre antes de fiarte de la capacidad.', wrap=True, b=True, fill=GRIS)
    ws.merge_cells('A3:I3')
    ws.row_dimensions[3].height = 36
    _bloque(ws, 'A. EL ASADOR Y EL DOMINGO', ENT_S, M_S, M_RS, SAL_S, S_S,
            ('carga', 'cap_h', 'tandas', 'cap_dom', 'sem', 'sem_ago', 'tarde_h', 'tarde_m', 'gas_pollo'))
    for k in ('sem', 'sem_ago'):
        _semaforo(ws, 'B%d' % S_S[k], SEM_OK, SEM_KO)
        ws.row_dimensions[S_S[k]].height = 32
    cel = lambda k: 'B%d' % M_S[k]
    _dv(ws, 'whole', [cel(k) for k in ('E', 'p', 'D')], f1='1', f2='1000000', op='between', titulo='Número no válido', msg='Escribe un número entero de 1 o más.')
    _dv(ws, 'decimal', [cel('t')], f1='0', op='greaterThan', msg='Escribe un número mayor que 0.')
    _dv(ws, 'decimal', [cel('c'), cel('kw'), cel('ekwh')], f1='0', op='greaterThanOrEqual', msg='Escribe un número mayor o igual que 0.')
    _dv(ws, 'whole', [cel('hh0'), cel('hh1')], f1='0', f2='23', op='between', titulo='Hora no válida', msg='Escribe una hora entera de 0 a 23.')
    _dv(ws, 'whole', [cel('mm0'), cel('mm1')], f1='0', f2='59', op='between', titulo='Minutos no válidos', msg='Escribe minutos enteros de 0 a 59.')
    _dv(ws, 'decimal', [cel('hp')], f1='0', f2='1', op='between', titulo='Porcentaje no válido', msg='Escribe un porcentaje entre 0 y 1 (0,8 = 80 %).')
    _dv(ws, 'decimal', [cel('ag')], f1='0', f2='3', op='between', titulo='Porcentaje no válido', msg='Escribe entre 0 y 3 (0,15 = 15 %).')
    _imprimir(ws, horizontal=True)


def hoja_escandallo(wb):
    ws = wb.create_sheet(HE)
    for col, w in (('A', 44), ('B', 14), ('C', 9)):
        ws.column_dimensions[col].width = w
    for col in PCOL:
        ws.column_dimensions[col].width = 15
    ws.column_dimensions[ORI].width = 46
    _f(ws, 'A1', 'ESCANDALLO DEL POLLO Y DE LOS COMBOS — ENTERO, MEDIO, CUARTO Y 6 COMBOS, EN GRAMOS', b=True, sz=14)
    _f(ws, 'A2', 'Escribe solo en las celdas VERDES (precios, gramos, nombres, fracciones de pollo y precios de carta). Importes sin IVA salvo el '
                 'precio de carta. Los precios son una estimación del kit, editable: pon los de tu proveedor. El coste del pollo viene de la '
                 'hoja «%s».' % HP, wrap=True)
    ws.merge_cells('A2:%s2' % ORI)
    ws.row_dimensions[2].height = 32
    for col, t in zip('ABD', ('DATOS GENERALES', 'Valor', 'Notas')):
        _f(ws, '%s4' % col, t, b=True, fill=GRIS)
    for k, rot, fmt, val, nota in ESC_IN:
        f = E[k]
        _f(ws, 'A%d' % f, rot)
        _f(ws, 'B%d' % f, val, fmt=fmt, fill=VERDE, locked=False)
        _nota(ws, f, nota, desde='D', hasta=ORI, chars=150.0)
    for col, t in (('A', 'INGREDIENTE (en gramos)'), ('B', '€/kg sin IVA'), ('C', 'Unidad')):
        _f(ws, '%s%d' % (col, E_CAB), t, b=True, fill=GRIS, wrap=True)
    _f(ws, '%s%d' % (ORI, E_CAB), 'ORIGEN DEL PRECIO', b=True, fill=GRIS, wrap=True)
    _f(ws, 'A%d' % E_PVP, 'Precio de carta con IVA (€)', b=True)
    _f(ws, 'A%d' % E_BEB, 'De ese precio, bebida que sale del local, con IVA (€)', wrap=True)
    _f(ws, 'A%d' % E_PAN, 'De ese precio, pan vendido suelto, con IVA (€)', wrap=True)
    ws.row_dimensions[E_BEB].height = 32
    ws.row_dimensions[E_PAN].height = 32
    for j, (nom, (pvp, beb, pan, _g, _u)) in enumerate(RECETAS.items()):
        col = PCOL[j]
        _f(ws, '%s%d' % (col, E_CAB), nom, b=True, fill=VERDE, wrap=True, locked=False)
        _f(ws, '%s%d' % (col, E_PVP), pvp, fmt=EUR, fill=VERDE, locked=False, b=True)
        _f(ws, '%s%d' % (col, E_BEB), beb, fmt=EUR, fill=VERDE, locked=False)
        _f(ws, '%s%d' % (col, E_PAN), pan, fmt=EUR, fill=VERDE, locked=False)
    ws.row_dimensions[E_CAB].height = 62
    for i, (nom, p, origen) in enumerate(ING_G):
        f = E_G0 + i
        _f(ws, 'A%d' % f, nom, fill=VERDE, locked=False)
        _f(ws, 'B%d' % f, p, fmt=EUR, fill=VERDE, locked=False)
        _f(ws, 'C%d' % f, 'g')
        for j, (_n, (_pvp, _b, _p2, gr, _ud)) in enumerate(RECETAS.items()):
            _f(ws, '%s%d' % (PCOL[j], f), float(gr.get(nom, 0)), fmt='#,##0.0;-#,##0.0;"-"', fill=VERDE, locked=False)
        _f(ws, '%s%d' % (ORI, f), origen, i=True)
    _f(ws, 'A%d' % E_UH, 'INGREDIENTE (en unidades)', b=True, fill=GRIS)
    _f(ws, 'B%d' % E_UH, '€/ud sin IVA', b=True, fill=GRIS)
    _f(ws, 'C%d' % E_UH, 'Unidad', b=True, fill=GRIS)
    for i, (nom, p, origen) in enumerate(ING_U):
        f = E_U0 + i
        if p is None:
            _f(ws, 'A%d' % f, nom)
            _f(ws, 'B%d' % f, "='%s'!$B$%d" % (HP, S_P['coste']), fmt=EUR3)
            _f(ws, 'C%d' % f, 'pollo')
        else:
            _f(ws, 'A%d' % f, nom, fill=VERDE, locked=False)
            _f(ws, 'B%d' % f, p, fmt=EUR3, fill=VERDE, locked=False)
            _f(ws, 'C%d' % f, 'ud')
        for j, (_n, (_pvp, _b, _p2, _g, ud)) in enumerate(RECETAS.items()):
            _f(ws, '%s%d' % (PCOL[j], f), float(ud.get(nom, 0)), fmt='#,##0.00;-#,##0.00;"-"', fill=VERDE, locked=False)
        _f(ws, '%s%d' % (ORI, f), origen, i=True)
    _f(ws, 'A%d' % E_R, 'RESULTADO POR PLATO (gramos y unidades en las columnas de arriba)', b=True, fill=GRIS)
    for j in range(NP):
        _f(ws, '%s%d' % (PCOL[j], E_R), '=%s%d' % (PCOL[j], E_CAB), b=True, fill=GRIS, wrap=True)
    ws.row_dimensions[E_R].height = 62
    kg = _kg_ext()
    gp = "'%s'!$B$%d" % (HP, S_P['g_entero'])
    for k, rot, fmt in RES:
        f = R[k]
        neg = k in ('c_tot', 'fc')
        _f(ws, 'A%d' % f, rot, b=neg)
        for j in range(NP):
            c = PCOL[j]
            rg = '%s$%d:%s$%d' % (c, E_G0, c, E_G1)
            ru = '%s$%d:%s$%d' % (c, E_U0 + 1, c, E_U1)
            ex = {
                'peso': 'SUM(%s)+%s$%d*%s' % (rg, c, E_U0, gp),
                'c_pollo': '%s$%d*$B$%d' % (c, E_U0, E_U0),
                'c_acomp': 'SUMPRODUCT(%s,$B$%d:$B$%d)/%s+SUMPRODUCT(%s,$B$%d:$B$%d)' % (rg, E_G0, E_G1, kg, ru, E_U0 + 1, E_U1),
                'c_tot': '%s%d+%s%d*(1+$B$%d)' % (c, R['c_pollo'], c, R['c_acomp'], E['merm']),
                'pvp_sin': '(%s%d-%s%d-%s%d)/(1+$B$%d)+%s%d/(1+$B$%d)+%s%d/(1+$B$%d)' % (
                    c, E_PVP, c, E_BEB, c, E_PAN, E['iva'], c, E_BEB, E['ivab'], c, E_PAN, E['ivap']),
                'fc': '%s%d/%s%d' % (c, R['c_tot'], c, R['pvp_sin']),
                'sem': 'IF(%s%d<=$B$%d,"%s",IF(%s%d<=$B$%d+$B$%d,"%s","%s"))' % (
                    c, R['fc'], E['obj'], FC_OK, c, R['fc'], E['obj'], E['tol'], FC_VIG, FC_ALTO),
                'margen': '%s%d-%s%d' % (c, R['pvp_sin'], c, R['c_tot']),
                'pvp_sug': '%s%d/$B$%d*(1+$B$%d)' % (c, R['c_tot'], E['obj'], E['iva']),
            }[k]
            _f(ws, '%s%d' % (c, f), '=' + ex, fmt=fmt, b=neg, wrap=(k == 'sem'))
    ws.row_dimensions[R['sem']].height = 32
    rng = '%s%d:%s%d' % (PCOL[0], R['sem'], PCOL[-1], R['sem'])
    top = '%s%d' % (PCOL[0], R['sem'])
    for txt, fill in ((FC_OK, OK), (FC_VIG, AMBAR), (FC_ALTO, ROJO)):
        ws.conditional_formatting.add(rng, FormulaRule(formula=['%s="%s"' % (top, txt)], fill=fill))
    # margen por canal de todos los platos
    _f(ws, 'A%d' % E_CH, 'MARGEN POR CANAL (todos los platos)', b=True, fill=GRIS)
    for j in range(NP):
        _f(ws, '%s%d' % (PCOL[j], E_CH), '=%s%d' % (PCOL[j], E_CAB), b=True, fill=GRIS, wrap=True)
    ws.row_dimensions[E_CH].height = 62
    for k, rot, fmt in CANAL:
        f = CH[k]
        _f(ws, 'A%d' % f, rot, b=k.startswith('plat') or k == 'rank', wrap=True)
        for j in range(NP):
            c = PCOL[j]
            ps = '%s%d' % (c, R['pvp_sin'])
            ex = {
                'sala_e': '%s%d' % (c, R['margen']),
                'sala_p': '%s%d/%s' % (c, CH['sala_e'], ps),
                'mostr_e': '%s%d-$B$%d' % (c, CH['sala_e'], E['env']),
                'mostr_p': '%s%d/%s' % (c, CH['mostr_e'], ps),
                'plat_pv': '%s*(1+$B$%d)' % (ps, E['rec']),
                'plat_e': '%s%d*(1-$B$%d)-%s%d-$B$%d' % (c, CH['plat_pv'], E['com'], c, R['c_tot'], E['env']),
                'plat_p': '%s%d/%s%d' % (c, CH['plat_e'], c, CH['plat_pv']),
                'rank': '1+COUNTIF($%s$%d:$%s$%d,">"&%s%d)' % (PCOL[0], CH['mostr_e'], PCOL[-1], CH['mostr_e'], c, CH['mostr_e']),
            }[k]
            _f(ws, '%s%d' % (c, f), '=' + ex, fmt=fmt, b=k.startswith('plat') or k == 'rank')
    ws.row_dimensions[CH['rank']].height = 32
    _f(ws, 'A%d' % E_FIN, 'El food cost de la tabla es el del plato servido en sala. En el mostrador y el encargo se descuenta el envase del '
                          'pedido; en la plataforma, además, se sube el precio con el recargo y se descuenta la comisión sobre ese precio. '
                          'AVISO FIJO: en la plataforma, vende el pollo entero con patatas asadas: las fritas llegan blandas. '
                          'Las bebidas sueltas no están en este escandallo.', i=True, wrap=True)
    ws.merge_cells('A%d:%s%d' % (E_FIN, ORI, E_FIN))
    ws.row_dimensions[E_FIN].height = 48
    # validaciones
    _dv(ws, 'decimal', ['B%d' % E[k] for k in ('iva', 'ivab', 'ivap', 'tol', 'merm', 'com', 'rec')], f1='0', f2='0.9', op='between',
        titulo='Porcentaje no válido', msg='Escribe un porcentaje entre 0 y 0,9 (0,1 = 10 %).')
    _dv(ws, 'decimal', ['B%d' % E['obj']], f1='0.01', f2='0.9', op='between', titulo='Objetivo no válido',
        msg='El objetivo no puede ser 0 (se divide entre él): escribe entre 0,01 y 0,9.')
    _dv(ws, 'decimal', ['B%d' % E['env']] + ['B%d' % f for f in range(E_G0, E_G1 + 1)] + ['B%d' % f for f in range(E_U0 + 1, E_U1 + 1)]
        + ['%s%d' % (c, f) for c in PCOL for f in [E_PVP, E_BEB, E_PAN] + list(range(E_G0, E_G1 + 1)) + list(range(E_U0, E_U1 + 1))],
        f1='0', op='greaterThanOrEqual', titulo='Importe no válido', msg='Escribe un número mayor o igual que 0.')
    _imprimir(ws, horizontal=True)


INSTRUCCIONES = [
    ('INSTRUCCIONES DE USO — Calculadora del pollo, el asador y el escandallo', True),
    ('Tres herramientas en un libro: la hoja «Pollo» te dice cuánto pesa un pollo asado, cuánto cuesta cada kilo y cada ración y cuántos pollos '
     'necesitas para un número de personas; la hoja «Asador» compara la capacidad de tus espadas con la punta del domingo y la de agosto; la hoja '
     '«Escandallo» calcula el coste, el food cost, el margen y el precio de carta del pollo entero, el medio, el cuarto y 6 combos, con el margen '
     'en sala, en el mostrador y en la plataforma. Solo se escribe en las celdas VERDES; lo demás se calcula.', False),
    ('Todos los valores por defecto son una estimación del kit, editable: son ejemplos coherentes con el caso del plan de negocio, no precios '
     'de nadie. Cámbialos por los de tu proveedor y por lo que midas en tu cocina antes de decidir precios.', False),
    ('HOJA «POLLO»', True),
    ('1. Cómo medir el rendimiento. Pesa 5 pollos crudos de la misma tanda, asa los 5 juntos y pésalos otra vez al sacarlos, ya escurridos. '
     'Divide el peso asado entre el peso en canal y usa la media: es tu rendimiento real. La referencia del USDA es 78 %, con un rango de '
     '68 a 84 %.', False),
    ('2. El resultado más útil es el coste por pollo entero y por kilo asado, que alimentan el escandallo, y los pollos que tienes que pedir '
     'para un encargo o un evento. Fíjate también en el coste de los sobrantes que se tiran: los pollos que quedan al cerrar y no se '
     'reaprovechan son dinero perdido.', False),
    ('3. Seguridad alimentaria. Mide la temperatura en el muslo con sonda: 74 °C o más es el criterio de este kit para dar un pollo por hecho, '
     'no una temperatura fijada por una norma. En la vitrina caliente, mantén 63 °C o más y fija en tu APPCC una hora límite propia. El '
     'sobrante que vayas a reaprovechar se enfría de 60 a 10 °C en menos de 2 horas (RD 1086/2020, art. 30) y se recalienta a 74 °C o más.', False),
    ('HOJA «ASADOR»', True),
    ('1. Escribe tus espadas, los pollos por espada y los minutos de asado. El tiempo de asado no tiene fuente publicada: cronometra una tanda '
     'desde que enciendes hasta que el muslo llega a 74 °C, con el calibre que compras. Suma los minutos de carga y descarga.', False),
    ('2. Escribe la hora del primer encendido y la hora límite en que quieres tener los pollos asados (la punta del domingo termina hacia '
     'las 15:00). La hoja cuenta las tandas completas que caben, la capacidad y si cubres la punta de un domingo normal y la de agosto.', False),
    ('3. El encendido más tardío es la hora a la que tienes que encender el asador para no quedarte corto en la punta. El gas por pollo usa la '
     'potencia media y el precio del gas que pongas: es una estimación, compruébala con tu factura o tu bombona.', False),
    ('HOJA «ESCANDALLO»', True),
    ('1. Escribe el IVA (10 % la comida, 21 % la bebida que sale del local, 4 % el pan vendido suelto), el food cost objetivo, el margen de '
     'vigilancia, la merma de las guarniciones, el envase, la comisión y el recargo de la plataforma.', False),
    ('2. En la tabla, cada fila es un ingrediente con su precio por kilo sin IVA. Los gramos de cada plato van en su columna. El pollo se mide '
     'en fracciones de pollo entero (1, 0,5, 0,25) y toma su coste de la hoja «Pollo». Si un precio de carta incluye bebida que sale del local o '
     'pan vendido suelto, escríbelo en las dos filas de debajo del precio para que el IVA se calcule bien.', False),
    ('3. Para cada plato ves el coste, el precio sin IVA, el food cost con semáforo, el margen y el precio de carta con IVA que necesitarías '
     'para llegar a tu objetivo. Abajo, el margen en sala, en el mostrador y el encargo y en la plataforma, y el puesto de cada plato por margen. '
     'El pollo entero suele quedar por encima del objetivo; los combos con guarniciones lo compensan.', False),
    ('PRECIOS, IVA Y FUENTES', True),
    ('El IVA de la comida preparada para consumir en el local, para llevar o a domicilio es el 10 % (Ley 37/1992, art. 91). La cerveza, el vino y los '
     'refrescos con azúcares o edulcorantes añadidos que salen del local llevan el 21 %; el pan común, el 4 %. El alioli, la ensaladilla y las cremas sin '
     'cocción se hacen con ovoproducto, a 8 °C como máximo y se consumen en 24 horas, con fecha y hora de elaboración (RD 1021/2022, art. 9).', False),
    ('Los precios y rendimientos salen del estudio de mercado del kit (octubre 2026) y de estimaciones del kit; la referencia de rendimiento '
     'viene del USDA. Ninguna cifra de esta calculadora es un dato publicado de tu negocio: sustitúyelas por las tuyas.', False),
]


def hoja_instrucciones(wb):
    ws = wb.create_sheet(HI)
    ws.column_dimensions['A'].width = 110
    f = 1
    for t, b in INSTRUCCIONES:
        _f(ws, 'A%d' % f, t, b=b, sz=13 if f == 1 else 11, wrap=True)
        if not b:
            ws.row_dimensions[f].height = max(16, 15 * math.ceil(len(t) / 120.0))
        f += 2 if b and f > 1 else 1
    f += 1
    for t in ('Para editar una celda que no esté en verde: Revisar → Desproteger hoja (no tiene contraseña).',
              'Más plantillas y kits del catálogo en aichef.pro/productos-digitales', motor.BIO_LINE,
              motor.VERSION_FMT.format(pid=PID)):
        _f(ws, 'A%d' % f, t, wrap=True)
        f += 1
    _imprimir(ws)


def generar(destino):
    wb = openpyxl.Workbook()
    hoja_pollo(wb)
    hoja_asador(wb)
    hoja_escandallo(wb)
    hoja_instrucciones(wb)
    pr = wb.properties
    pr.title = '%s · v%s' % (TITULO, motor.VERSION)
    pr.subject = 'Plan de Negocio: Pollería y Asador de Pollos'
    pr.description = 'aichef.pro/%s' % PID
    pr.keywords = 'plan negocio pollería, calculadora asador de pollos, escandallo pollo asado, AI Chef Pro'
    pr.creator = pr.lastModifiedBy = motor.CREATOR
    os.makedirs(os.path.dirname(destino), exist_ok=True)
    wb.save(destino)
    sys.path.insert(0, PD)
    import inject_cache as IC                                                      # noqa: E402
    nf, inj, fail = IC.inject(destino)
    print('%s: formulas=%d cache_inyectado=%d fallos_pycel=%d' % (os.path.basename(destino), nf, inj, fail))
    return fail


def valores_expuestos(ruta):
    """Lee de la CACHÉ del xlsx los 5 valores expuestos: {nombre: valor}."""
    wv = openpyxl.load_workbook(ruta, data_only=True)
    return OrderedDict((n, wv[h][c].value) for n, (h, c) in EXPUESTOS.items())


# ---------------------------------------------------------------------------------------------------------------------
# Autotest
# ---------------------------------------------------------------------------------------------------------------------
PROHIBIDOS = [r'pizz', r'napol[ei]t', r'taquer', r'trompo', r'\btacos?\b', r'tortilla', r'\bcomal\b', r'al pastor', r'gringa', r'nixtamal',
              r'mezcal', r'tequila', r'aguas? frescas?', r'carnitas', r'suadero', r'cochinita', r'guacamole', r'salsa macha',
              r'plan-negocio-taqueria', r'horchata', r'mantec', r'cubeta', r'levadura', r'mozzarella', r'sin gluten', r'auténtic',
              r'mexican', r'broaster', r'digital obligatorio', r'sales tax|service charge', r'642\.5', r'chimenea',
              r'obligatori[oa]s?[^.]{0,40}74|74[^.]{0,40}obligatori']
CELDA = {}                                                                         # clave de entrada → (hoja, celda)
for _k, _f2 in M_P.items():
    CELDA[_k] = (HP, 'B%d' % _f2)
for _k, _f2 in M_S.items():
    CELDA[_k] = (HA, 'B%d' % _f2)
for _k, _f2 in E.items():
    CELDA[_k] = (HE, 'B%d' % _f2)
for _i in range(NG):
    CELDA['p%d' % _i] = (HE, 'B%d' % (E_G0 + _i))
for _i in range(1, NU):
    CELDA['up%d' % _i] = (HE, 'B%d' % (E_U0 + _i))
for _j in range(NP):
    CELDA['pvp%d' % _j] = (HE, '%s%d' % (PCOL[_j], E_PVP))
    CELDA['beb%d' % _j] = (HE, '%s%d' % (PCOL[_j], E_BEB))
    CELDA['pan%d' % _j] = (HE, '%s%d' % (PCOL[_j], E_PAN))
    for _i in range(NG):
        CELDA['g%d_%d' % (_i, _j)] = (HE, '%s%d' % (PCOL[_j], E_G0 + _i))
    for _i in range(NU):
        CELDA['u%d_%d' % (_i, _j)] = (HE, '%s%d' % (PCOL[_j], E_U0 + _i))


def celdas_salida(r):
    """(hoja, celda, valor esperado, etiqueta) de todas las salidas."""
    out = []
    for k in S_P:
        out.append((HP, 'B%d' % S_P[k], r['P'][k], 'P.' + k))
    for k in S_S:
        out.append((HA, 'B%d' % S_S[k], r['S'][k], 'S.' + k))
    for i in range(NG):
        out.append((HE, 'B%d' % (E_G0 + i), r['pg'][i], 'precio%d' % i))
    for i in range(NU):
        out.append((HE, 'B%d' % (E_U0 + i), r['pu'][i], 'precio_ud%d' % i))
    for j in range(NP):
        for k in list(R) + list(CH):
            out.append((HE, '%s%d' % (PCOL[j], R[k] if k in R else CH[k]), r['p%d' % j][k], '%s[%d]' % (k, j)))
    return out


def _distinto(x, v):
    if isinstance(v, str):
        return x != v
    return not isinstance(x, (int, float)) or abs(float(x) - v) > 0.01


def _checks_logica(nombre, jj, err):
    r = calcular(jj)
    p, s = r['P'], r['S']
    if abs(p['asado'] - jj['P'] * jj['r']) > 1e-9:
        err.append('%s: asado ≠ P × r' % nombre)
    if abs(p['c_medio'] * jj['M2'] - p['coste']) > 0.01 or abs(p['c_cuarto'] * jj['M4'] - p['coste']) > 0.01:
        err.append('%s: ½ × 2 o ¼ × 4 ≠ pollo entero' % nombre)
    if abs(p['g_medio'] * jj['M2'] - p['asado'] * jj['KG']) > 0.01 or abs(p['g_cuarto'] * jj['M4'] - p['asado'] * jj['KG']) > 0.01:
        err.append('%s: raciones × peso ≠ peso asado' % nombre)
    if abs(p['coste_kg'] * p['asado'] - p['coste']) > 0.01:
        err.append('%s: coste por kg × peso asado ≠ coste del pollo' % nombre)
    if p['pollos_N'] * jj['pp'] < jj['Np'] or (p['pollos_N'] - 1) * jj['pp'] >= jj['Np']:
        err.append('%s: pollos para N personas no es el techo de N ÷ personas por pollo' % nombre)
    if abs(s['cap_dom'] - jj['E'] * jj['p'] * s['tandas']) > 1e-9:
        err.append('%s: capacidad ≠ espadas × pollos × tandas' % nombre)
    if s['cargas_dom'] * s['carga'] < jj['D'] or (s['cargas_dom'] - 1) * s['carga'] >= jj['D']:
        err.append('%s: tandas del domingo no son el techo de D ÷ (E × p)' % nombre)
    if s['min_disp'] >= 0 and not (s['tandas'] * s['ciclo'] <= s['min_disp'] + 1e-9 < (s['tandas'] + 1) * s['ciclo'] + 1e-9):
        err.append('%s: tandas × ciclo no encaja en los minutos disponibles' % nombre)
    if abs(s['cap_h'] * s['ciclo'] - s['carga'] * jj['MH']) > 1e-6:
        err.append('%s: capacidad por hora × ciclo ≠ pollos por tanda × 60' % nombre)
    if abs(s['tarde_h'] * jj['MH'] + s['tarde_m'] - s['tarde_min']) > 1e-9:
        err.append('%s: encendido más tardío hora y minutos no suman' % nombre)
    for j in range(NP):
        q = r['p%d' % j]
        frac = jj['u0_%d' % j]
        if abs(q['c_pollo'] - frac * p['coste']) > 0.01:
            err.append('%s: coste del pollo de la columna %d ≠ fracción × coste del pollo' % (nombre, j))
        if not (q['plat_e'] < q['mostr_e'] < q['sala_e'] and q['plat_p'] < q['mostr_p'] < q['sala_p']):
            err.append('%s: margen plataforma < mostrador < sala no se cumple en la columna %d' % (nombre, j))


def autotest(destino):
    from pycel import ExcelCompiler
    err, n = [], 0
    wv = openpyxl.load_workbook(destino, data_only=True)
    wf = openpyxl.load_workbook(destino)
    base = valores_defecto()
    j2 = dict(base, P=1.6, Pk=3.25, S=0.20, r=0.72, pp=4, Np=75, Ns=3.0, a=0.2, E=8, p=6, t=85.0, c=10.0, hh0=11, mm0=15, hh1=15, mm1=30,
              D=140, hp=0.7, ag=0.25, kw=56.0, ekwh=0.12, iva=0.10, ivab=0.21, ivap=0.04, obj=0.30, tol=0.04, merm=0.05, env=0.7, com=0.25,
              rec=0.10, p0=2.8, p4=3.5, up1=0.2, up2=0.6, pvp0=13.5, pvp8=22.5, beb8=2.0, pan7=1.0, g0_3=400.0, g2_4=180.0, u0_1=0.5, u2_8=2.0)
    j3 = dict(base, P=1.3, Pk=2.9, r=0.8, pp=2, Np=10, Ns=0.0, E=6, p=4, t=70.0, c=20.0, hh0=9, mm0=0, hh1=14, mm1=45, D=60, hp=0.9, ag=0.0,
              kw=40.0, ekwh=0.08, obj=0.4, tol=0.1, merm=0.0, env=0.4, com=0.35, rec=0.2, p1=0.9, p6=5.0, up1=0.15, pvp3=17.5, beb8=0.0,
              pan8=0.9, g1_6=500.0, g5_7=60.0, u0_7=1.5, u1_5=2.0)
    # 1. caché y pycel del juego por defecto; pycel de los otros dos con set_value
    xl = ExcelCompiler(filename=destino)
    r0 = calcular(base)
    for hoja, c, v, lab in celdas_salida(r0):
        n += 1
        for et, x in (('pycel', xl.evaluate("'%s'!%s" % (hoja, c))), ('caché', wv[hoja][c].value)):
            if _distinto(x, v):
                err.append('defecto · %s!%s (%s): %s %r ≠ esperado %r' % (hoja, c, lab, et, x, v))
    for nombre, jj in (('juego 2', j2), ('juego 3', j3)):
        r = calcular(jj)
        xl2 = ExcelCompiler(filename=destino)
        for hoja, c, v, lab in celdas_salida(r0):
            xl2.evaluate("'%s'!%s" % (hoja, c))
        for k, v in jj.items():
            if v != base[k]:
                h, c = CELDA[k]
                xl2.set_value("'%s'!%s" % (h, c), v)
        for hoja, c, v, lab in celdas_salida(r):
            n += 1
            x = xl2.evaluate("'%s'!%s" % (hoja, c))
            if _distinto(x, v):
                err.append('%s · %s!%s (%s): pycel %r ≠ esperado %r' % (nombre, hoja, c, lab, x, v))
    # 2. identidades de la lógica en los tres juegos
    for nombre, jj in (('defecto', base), ('juego 2', j2), ('juego 3', j3)):
        n += 12
        _checks_logica(nombre, jj, err)
    # identidades sobre la hoja viva (pycel, valores por defecto)
    ev = lambda h, c: xl.evaluate("'%s'!%s" % (h, c))
    asado, pb, rr = ev(HP, 'B%d' % S_P['asado']), ev(HP, 'B%d' % M_P['P']), ev(HP, 'B%d' % M_P['r'])
    gm, gc, ge = (ev(HP, 'B%d' % S_P[k]) for k in ('g_medio', 'g_cuarto', 'g_entero'))
    cap, ca, ta = (ev(HA, 'B%d' % S_S[k]) for k in ('cap_dom', 'carga', 'tandas'))
    n += 3
    if abs(asado - pb * rr) > 1e-6 or abs(gm * 2 - ge) > 0.01 or abs(gc * 4 - ge) > 0.01:
        err.append('hoja Pollo: asado ≠ P × r o ½ × 2 / ¼ × 4 ≠ entero')
    if abs(cap - ca * ta) > 1e-9 or ca != ev(HA, 'B%d' % M_S['E']) * ev(HA, 'B%d' % M_S['p']):
        err.append('hoja Asador: capacidad ≠ espadas × pollos por espada × tandas')
    # 3. valores con sentido
    ent = r0['p%d' % J_ENTERO]
    if not 0.35 <= ent['fc'] <= 0.42:
        err.append('food cost del pollo entero %.1f %% fuera de 35-42 %%' % (ent['fc'] * 100))
    if abs(r0['P']['coste'] - 4.5525) > 1e-6 or abs(r0['P']['asado'] - 1.102) > 1e-6 or abs(r0['P']['coste_kg'] - 4.1311) > 0.001:
        err.append('coste %.4f / asado %.4f / €/kg asado %.4f fuera de 4,5525 / 1,102 / 4,131' % (r0['P']['coste'], r0['P']['asado'], r0['P']['coste_kg']))
    if r0['S']['carga'] != 60 or r0['S']['tandas'] != 3 or r0['S']['cap_dom'] != 180 or r0['S']['cargas_dom'] != 2 or r0['S']['cargas_ago'] != 3:
        err.append('asador por defecto: 60 por tanda · 3 tandas · 180 · 2 tandas el domingo · 3 en agosto (salió %r)'
                   % [r0['S'][k] for k in ('carga', 'tandas', 'cap_dom', 'cargas_dom', 'cargas_ago')])
    if r0['S']['sem'] != SEM_OK or r0['S']['sem_ago'] != SEM_OK or abs(r0['S']['cap_h'] - 40) > 1e-9:
        err.append('asador por defecto: semáforos o capacidad por hora (40) inesperados')
    if calcular(dict(base, D=300))['S']['sem'] != SEM_KO:
        err.append('el semáforo de la punta no pasa a «No llegas» con 300 pollos')
    medio, cuarto = r0['p1'], r0['p2']
    if abs(medio['c_pollo'] * 2 - ent['c_pollo']) > 0.01 or abs(cuarto['c_pollo'] * 4 - ent['c_pollo']) > 0.01:
        err.append('½ × 2 o ¼ × 4 ≠ pollo entero en el escandallo')
    if not all(0.15 <= r0['p%d' % j]['fc'] <= 0.50 for j in range(NP)):
        err.append('algún food cost por defecto fuera de 15-50 %%: %r' % [round(r0['p%d' % j]['fc'], 3) for j in range(NP)])
    if sorted(r0['p%d' % j]['rank'] for j in range(NP)) != list(range(1, NP + 1)):
        err.append('los puestos por margen no son 1..%d' % NP)
    # 4. los 5 valores expuestos: caché frente a calcular()
    esperados = {'POLLOS_POR_TANDA': r0['S']['carga'], 'CAPACIDAD_DOMINGO': r0['S']['cap_dom'], 'RENDIMIENTO_POLLO': r0['P']['rend'],
                 'COSTE_POLLO_ENTERO': r0['P']['coste'], 'FC_POLLO_ENTERO': ent['fc']}
    expuestos = valores_expuestos(destino)
    for nombre, e_ in esperados.items():
        n += 1
        if _distinto(expuestos[nombre], e_):
            err.append('expuesto %s: caché %r ≠ %r' % (nombre, expuestos[nombre], e_))
    # 5. formato, protección, DV, fórmulas, textos
    if motor.gate_formatos(wf, FICHERO):
        err.append('gate_formatos: %r' % motor.gate_formatos(wf, FICHERO)[:3])
    dvh = dict((w.title, dict((str(rng), dv) for dv in w.data_validations.dataValidation for rng in dv.sqref.ranges))
               for w in (wf[HP], wf[HA], wf[HE]))
    for hoja, k, tipo in ((HP, 'P', 'decimal'), (HP, 'r', 'decimal'), (HP, 'Np', 'whole'), (HA, 'E', 'whole'), (HA, 't', 'decimal'),
                          (HA, 'hh0', 'whole'), (HA, 'mm1', 'whole'), (HA, 'D', 'whole'), (HE, 'obj', 'decimal'), (HE, 'iva', 'decimal'),
                          (HE, 'com', 'decimal')):
        c = CELDA[k][1]
        dv = dvh[hoja].get(c)
        if dv is None or dv.type != tipo:
            err.append('DV ausente o de otro tipo en %s!%s (esperada %s)' % (hoja, c, tipo))
    textos, todo_celdas = [], []
    for w in wf.worksheets:
        if w.page_setup.paperSize != 9 or not w.protection.sheet:
            err.append('%s: A4 y protección' % w.title)
        for row in w.iter_rows():
            for c in row:
                v = c.value
                if v is not None:
                    todo_celdas.append(str(v))
                if isinstance(v, str) and c.data_type != 'f':
                    textos.append(v)
                if c.data_type == 'f':
                    if wv[w.title][c.coordinate].value is None:
                        err.append('%s!%s: fórmula sin caché' % (w.title, c.coordinate))
                    sin_refs = re.sub(r'"[^"]*"', '', re.sub(r"'[^']+'!|\$?[A-Z]{1,2}\$?\d+", '', str(c.value)))
                    for num in re.findall(r'(?<![A-Za-z])\d+\.?\d*', sin_refs):
                        if num not in ('0', '1'):
                            err.append('%s!%s: número tecleado %s en %s' % (w.title, c.coordinate, num, c.value))
                if c.fill.fgColor.rgb in ('00E8F5E9', 'FFE8F5E9') and c.protection.locked:
                    err.append('%s!%s: verde bloqueada' % (w.title, c.coordinate))
                if c.fill.fgColor.rgb in ('00E8F5E9', 'FFE8F5E9') and c.data_type == 'f':
                    err.append('%s!%s: verde con fórmula' % (w.title, c.coordinate))
    for k in ('KG', 'M2', 'M4', 'ref78'):
        c = wf[HP][CELDA[k][1]]
        if c.fill.fgColor.rgb not in ('00F2F2F2', 'FFF2F2F2') or c.protection.locked is not True:
            err.append('constante gris %s mal marcada o desbloqueada' % k)
    c = wf[HA][CELDA['MH'][1]]
    if c.fill.fgColor.rgb not in ('00F2F2F2', 'FFF2F2F2') or c.protection.locked is not True:
        err.append('constante gris MH mal marcada o desbloqueada')
    todo = '\n'.join(textos)
    barrido = '\n'.join(todo_celdas)
    for rx in PROHIBIDOS:
        if re.search(rx, barrido, re.I):
            err.append('prohibido /%s/' % rx)
    if re.search(r'[^\x00-ɏ‐-‧€−→·]', barrido):
        err.append('caracteres no latinos o raros: %r' % sorted(set(re.findall(r'[^\x00-ɏ‐-‧€−→·]', barrido))))
    if 'Estimación del kit, editable' not in barrido and 'estimación del kit, editable' not in barrido:
        err.append('falta el rótulo «estimación del kit, editable»')
    if motor.VERSION_FMT.format(pid=PID) not in todo or motor.BIO_LINE not in todo:
        err.append('falta la línea de versión o la bio')
    if '1.0' not in motor.VERSION_FMT or 'noviembre 2026' not in motor.VERSION_FMT:
        err.append('versión no es 1.0 · noviembre 2026: %s' % motor.VERSION_FMT)
    if wf.sheetnames != [HP, HA, HE, HI]:
        err.append('pestañas %r' % wf.sheetnames)
    if NG + NU < 10:
        err.append('menos de 10 ingredientes')
    p = wf.properties
    if p.creator != 'AI Chef Pro' or p.description != 'aichef.pro/%s' % PID or 'AI Chef Pro' not in (p.keywords or ''):
        err.append('docProps: creator=%r description=%r keywords=%r' % (p.creator, p.description, p.keywords))
    for x in err:
        print('ERROR', x)
    print('autotest calculadora pollería ES: %d comprobaciones (pycel + caché por defecto; pycel en 2 juegos más; identidades) · %s'
          % (n, 'VERDE' if not err else 'ROJO (%d errores)' % len(err)))
    P_, S_ = r0['P'], r0['S']
    print('  pollo: %.2f kg → %.3f kg asado (rendimiento %.1f %%) · coste %.4f € · %.3f €/kg asado · ½ %.0f g %.3f € · ¼ %.0f g %.3f €'
          % (base['P'], P_['asado'], P_['rend'] * 100, P_['coste'], P_['coste_kg'], P_['g_medio'], P_['c_medio'], P_['g_cuarto'], P_['c_cuarto']))
    print('  asador: %d por tanda · %.0f por hora · %d tandas antes de la hora límite · capacidad %d · punta %.1f · %s · agosto %s · '
          'encendido más tardío %02d:%02d · gas %.3f €/pollo' % (S_['carga'], S_['cap_h'], S_['tandas'], S_['cap_dom'], S_['dem_punta'],
                                                               S_['sem'], S_['sem_ago'], S_['tarde_h'], S_['tarde_m'], S_['gas_pollo']))
    for nombre, (h, c) in EXPUESTOS.items():
        print('  %-20s %s!%-5s %s' % (nombre, h, c, expuestos[nombre]))
    for j, nom in enumerate(RECETAS):
        q = r0['p%d' % j]
        print('  %-46s coste %.3f € · food cost %5.1f %% · %s · margen sala %.2f / mostrador %.2f / plataforma %.2f € · puesto %d'
              % (nom, q['c_tot'], q['fc'] * 100, q['sem'], q['sala_e'], q['mostr_e'], q['plat_e'], q['rank']))
    return 1 if err else 0


if __name__ == '__main__':
    ap = argparse.ArgumentParser()
    ap.add_argument('--salida', required=True, help='carpeta de salida')
    ap.add_argument('--autotest', action='store_true')
    a = ap.parse_args()
    destino = os.path.join(a.salida, FICHERO)
    if not a.autotest and generar(destino):
        sys.exit('ERROR: fallos de pycel en inject_cache')
    sys.exit(autotest(destino))
