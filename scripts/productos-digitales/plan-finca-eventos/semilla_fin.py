#!/usr/bin/env python3
"""
semilla_fin.py — Plan de Negocio: Finca de Eventos y Bodas (ES 1.0; SPEC D3, D6, D7): prepara el ORIGEN que
`planes-v2_0/main.py` convierte en el plan financiero y el checklist de la finca. Copia adaptada de
`plan-heladeria/semilla_hel.py` (F1-camino §3), sobre el MISMO origen.

Copia los dos xlsx de la Panadería ES 2.2 publicados (`dl/plan-negocio-panaderia/`, SOLO LECTURA) con los nombres de
fichero de la finca y escribe lo que el módulo de contenido no alcanza (F1-camino §1):
  · plan: valores de `0. Supuestos` (grupo_a conserva lo verde que trae el libro), títulos de hoja (Inversión A1/A2,
    PyG A2, Punto Equilibrio A2), las 19 partidas de inversión (filas 7-25: rótulo, importe, nota, ¿IVA?), los 10
    costes fijos propios del P&L (filas 26-35: rótulo e importe; la nota la pone el contenido), `PyG!G13` por valor
    (IVA soportado del coste directo), vacía `PyG!H4:H11` (la línea partida del pan común) y los docProps;
  · checklist: título, pestañas, cabeceras de fase y las tareas de `datos_checklist.py` con el pie «Tareas
    completadas:» debajo; `Instrucciones!A3` con el pid; DV, formato condicional y contadores los rehace el motor.
Corre en el VPS (openpyxl). Después: `main_fin.py --producto plan-negocio-finca-eventos-bodas --dry-run --origen <salida>`.

    /root/venv-guias/bin/python semilla_fin.py --salida /root/tmp-fin-f2/motor/origen

Sesión Claude Code, 25-oct-2026.
"""
import argparse
import copy
import importlib
import os
import shutil
import sys

import openpyxl
from openpyxl.styles import PatternFill

import datos_checklist as DC

AQUI = os.path.dirname(os.path.abspath(__file__))
REPO = os.path.abspath(os.path.join(AQUI, '..', '..', '..'))
PAN = os.path.join(REPO, 'astro-site', 'public', 'dl', 'plan-negocio-panaderia')
PLAN_PAN, CHECK_PAN = 'plan-financiero-panaderia.xlsx', 'checklist-apertura-panaderia.xlsx'
PLAN, CHECK = 'plan-financiero-finca-eventos-bodas.xlsx', 'checklist-apertura-finca-eventos-bodas.xlsx'
PID = 'plan-negocio-finca-eventos-bodas'

TITULOS_PLAN = {
    ('Inversión Inicial', 'A1'): 'PLAN DE NEGOCIO: FINCA DE EVENTOS Y BODAS',
    ('Inversión Inicial', 'A2'): ('Inversión inicial estimada — finca arrendada con edificaciones existentes y catering '
                                  'exclusivo concesionado'),
    ('PyG 3 Años', 'A2'): 'Escenario realista — finca de eventos y bodas',
    ('Punto Equilibrio', 'A2'): 'Punto de equilibrio — finca de eventos y bodas',
}

#: Filas 7-25 de 'Inversión Inicial' (molde de la panadería): (rótulo, importe, nota, ¿lleva IVA?). Research §3: las 15
#: partidas de 242.000 €, desdobladas en 19 (SPEC D7) y ajustadas para que, con el 8 % de imprevistos sobre la obra
#: (fila 28, por fórmula), la inversión quede en ≈ 242.000 €: 230.000 € de partidas (obra 141.000 · maquinaria y
#: mobiliario 66.000 · no amortizable 23.000) + 11.280 € de imprevistos = 241.280 €. Sin palabras de recuento en el
#: rótulo (motor.RX_RECUENTO) ni «local» o «sala» (VOCABULARIO).
INVERSION = [
    ('Proyecto técnico y dirección de obra', 7000,
     'Proyecto de adecuación y de la actividad, dirección de obra y certificado final. Es la base de la calificación '
     'o autorización urbanística y de la licencia', 'Sí'),
    ('Calificación urbanística, licencia, tasas y prestación compensatoria o canon', 10000,
     'Lo que cobra la administración en suelo rústico, distinto en cada comunidad: en Andalucía, prestación '
     'compensatoria del 10 % de las obras sin maquinaria (LISTA, art. 22.5); en la C. Valenciana, canon del 2 % '
     '(TRLOTUP, art. 221). Ajusta esta partida a la tuya: compruébalo en tu ayuntamiento y en tu comunidad', 'No'),
    ('Adecuación del salón: suelos, cerramientos, carpintería y pintura', 30000,
     'Sobre edificaciones EXISTENTES: en suelo rústico es lo que exigen Madrid y Cataluña y la vía rápida de la C. '
     'Valenciana. Pide dos presupuestos sobre el proyecto', 'Sí'),
    ('Acústica y aislamiento del salón', 9000,
     'Para que el baile final pueda ir dentro y cumplir los 45 dBA de noche al exterior en zona residencial '
     '(RD 1367/2007, anexo III)', 'Sí'),
    ('Carpa estructural o pérgola para ceremonia y cóctel', 28000,
     'Es el plan B de lluvia: inversión, no extra. Una carpa es una instalación desmontable: para el plan de '
     'autoprotección estatal tiene el mismo umbral que un edificio cerrado, 1.000 personas (Orden INT/1068/2026)',
     'Sí'),
    ('Aseos: dos bloques y uno adaptado', 22000,
     'Con un aseo adaptado e itinerario accesible: compruébalo con tu técnico sobre el Código Técnico', 'Sí'),
    ('Instalación eléctrica: cuadro, alumbrado exterior y de emergencia', 15000,
     'Potencia para el salón, la carpa, el office y el exterior, con su boletín. Pregunta la potencia disponible '
     'ANTES de firmar: una acometida nueva en el campo tarda', 'Sí'),
    ('Grupo electrógeno', 10000,
     'Respaldo si la red falla o no llega la potencia: un corte de luz en mitad de una boda no tiene arreglo', 'Sí'),
    ('Climatización del salón', 17000, 'Frío y calor para usar el salón todo el año, también en las bodas de agosto',
     'Sí'),
    ('Office del catering: cámara, fregaderos, mesas inox, agua y desagüe', 16000,
     'El espacio donde el catering emplata y friega. La cocina de producción la pone él: por eso el plan no lleva '
     'cocina. Lo inscribe en el registro sanitario quien lo usa; compruébalo en tu comunidad', 'Sí'),
    ('Mobiliario de ceremonia y banquete: mesas, sillas y mantelería base', 18000,
     'Mesas redondas e imperiales, 250 sillas y mantelería base. La decoración especial la contratan los novios',
     'Sí'),
    ('Zona chill-out y decoración base', 4000, 'Sofás, iluminación ambiental y elementos fijos de decoración', 'Sí'),
    ('Sonido, limitador-registrador y proyector', 6000,
     'El limitador-registrador es tu seguro frente a las denuncias por ruido: muchas ordenanzas lo exigen. '
     'Compruébalo en la tuya', 'Sí'),
    ('Aparcamiento y caminos: zahorra, señalización e iluminación', 10000,
     'Calcula una plaza por cada 2,5-3 invitados (estimación del kit) y un acceso que aguante la lluvia', 'Sí'),
    ('Jardinería, paisajismo y riego inicial', 8000,
     'Zona de ceremonia, cóctel y fotos. El mantenimiento anual va en los costes fijos', 'Sí'),
    ('Protección contra incendios, señalización de evacuación y accesibilidad', 7000,
     'Extintores, alumbrado y señales de evacuación y los ajustes de accesibilidad que pida la licencia', 'Sí'),
    ('Web, fotos, vídeo y marca', 6000,
     'Identidad, web con calendario de fechas, sesión de fotos y vídeo de la finca montada', 'Sí'),
    ('Constitución de la SL', 1000, 'Notaría, registro mercantil y gestoría', 'Sí'),
    ('Lanzamiento: feria nupcial, puertas abiertas y escaparate premium', 6000,
     'La primera feria de novios, la jornada de puertas abiertas con el catering y el escaparate de pago en un '
     'portal de bodas durante el primer año', 'Sí'),
]
assert len(INVERSION) == 19 and sum(x[1] for x in INVERSION) == 230000, sum(x[1] for x in INVERSION)

#: Filas 26-35 de 'PyG 3 Años' (costes fijos propios). Research §3: 43.000 €/año en 10 partidas, una por fila del molde
#: (no se suprimen filas: desplazaría el P&L). SGAE y AGEDI-AIE van por evento en el coste directo (SPEC D6). Sin
#: palabras de recuento en el rótulo. Las claves de `contenido….FIJOS` son estos rótulos normalizados.
FIJOS = {
    26: ('Marketing: portales de bodas, Google Ads y redes', 12000),
    27: ('Ferias nupciales y jornadas de puertas abiertas', 6000),
    28: ('Jardinería y mantenimiento externo (poda, césped y riego)', 8000),
    29: ('Reparaciones, pintura y reposición de mobiliario', 6000),
    30: ('Gestoría laboral, fiscal y contable', 3000),
    31: ('Software de reservas, CRM, web y teléfono', 1800),
    32: ('Prevención de riesgos laborales y vigilancia de la salud', 500),
    33: ('Tasas municipales e IBI repercutido', 2500),
    34: ('Revisiones de baja tensión, protección contra incendios y grupo electrógeno', 1200),
    35: ('Limpieza general, desinsectación y desratización (DDD)', 2000),
}
assert sum(x[1] for x in FIJOS.values()) == 43000


def _contenido():
    sys.path.insert(0, os.path.join(os.path.dirname(AQUI), 'planes-v2_0'))
    import grupo_a
    return grupo_a, importlib.import_module('contenido_plan_negocio_finca_eventos_bodas.a')


def _supuestos(wb):
    """Los valores de `0. Supuestos` del caso finca: `grupo_a` CONSERVA lo que el libro ya trae en las celdas verdes y
    sólo usa el valor del contenido en una celda vacía; sin este paso el caso seguiría siendo el de la panadería."""
    grupo_a, contenido = _contenido()
    coords = dict((t[0], t[1]) for t in grupo_a.SUPUESTOS_BASE)
    ws = wb['0. Supuestos']
    escritas = 0
    for clave, t in contenido.SUPUESTOS.items():
        if t[2] is None:
            continue
        assert clave in coords, clave
        ws[coords[clave]] = t[2]
        escritas += 1
    return escritas


def _docprops(wb):
    """El motor no reescribe `description` ni `keywords` y heredaba los de la panadería."""
    wb.properties.description = 'aichef.pro/' + PID
    wb.properties.keywords = 'plan de negocio finca de eventos y bodas, AI Chef Pro'


def plan(salida):
    destino = os.path.join(salida, PLAN)
    shutil.copy2(os.path.join(PAN, PLAN_PAN), destino)
    wb = openpyxl.load_workbook(destino)
    print('  supuestos escritos:', _supuestos(wb))
    _docprops(wb)
    for (hoja, coord), texto in TITULOS_PLAN.items():
        wb[hoja][coord] = texto
    ws = wb['Inversión Inicial']
    for i, (rot, imp, nota, iva) in enumerate(INVERSION):
        r = 7 + i
        assert not str(ws.cell(r, 1).value).lower().startswith(('fianza', 'alquiler', 'imprevistos')), r
        ws.cell(r, 1).value, ws.cell(r, 2).value, ws.cell(r, 4).value, ws.cell(r, 5).value = rot, imp, nota, iva
    assert str(ws['A26'].value).lower().startswith('fianza'), ws['A26'].value
    assert str(ws['A27'].value).lower().startswith('alquiler'), ws['A27'].value
    assert str(ws['A28'].value).lower().startswith('imprevistos'), ws['A28'].value
    ws = wb['PyG 3 Años']
    assert str(ws['A25'].value).startswith('Seguros') and str(ws['A36'].value).startswith('Amortización')
    for r, (rot, imp) in FIJOS.items():
        ws.cell(r, 1).value, ws.cell(r, 2).value = rot, imp
        ws.cell(r, 6).value = None          # la nota la pone el contenido (FIJOS)
    # IVA soportado del coste directo (fila «Coste de mercancía — comida»): en la panadería era una FÓRMULA (peso del
    # pan común × 4 %); aquí es la celda verde de `IVA_COMPRAS` (21 %), escrita por valor para que el motor no la
    # cuente como una fórmula perdida. La línea partida del pan (H4:H11) se vacía: grupo_a no la reescribe.
    _g, contenido = _contenido()
    assert str(ws['A13'].value).startswith('Coste de mercancía — comida'), ws['A13'].value
    assert str(ws['G13'].value).startswith('='), ws['G13'].value
    ws['G13'] = contenido.IVA_COMPRAS['comida']
    # SPEC D4: las dos líneas de venta al 21 % (sexto campo de LINEAS_INGRESO → celda G verde). En la panadería G10 y
    # G11 eran fórmulas: se escriben por valor por la misma razón que G13.
    for r, linea in ((10, contenido.LINEAS_INGRESO[0]), (11, contenido.LINEAS_INGRESO[1])):
        assert str(ws.cell(r, 7).value).startswith('='), (r, ws.cell(r, 7).value)
        ws.cell(r, 7).value = linea[5]
    assert 'pan' in str(ws['H4'].value).lower(), ws['H4'].value
    for r in range(4, 12):
        c = ws.cell(r, 8)
        c.value = None
        c.fill = PatternFill(fill_type=None)
    wb.save(destino)
    return destino


def _copiar_estilo(origen, destino):
    destino.font, destino.border = copy.copy(origen.font), copy.copy(origen.border)
    destino.fill, destino.alignment = copy.copy(origen.fill), copy.copy(origen.alignment)
    destino.number_format, destino.protection = origen.number_format, copy.copy(origen.protection)


def checklist(salida):
    destino = os.path.join(salida, CHECK)
    shutil.copy2(os.path.join(PAN, CHECK_PAN), destino)
    wb = openpyxl.load_workbook(destino)
    fases = [ws for ws in wb.worksheets if ws.title[:2] in ('F1', 'F2', 'F3', 'F4', 'F5', 'F6')]
    assert len(fases) == 6
    for ws, (titulo, cabecera, tareas) in zip(fases, DC.FASES):
        assert len(titulo) <= 31, titulo
        ws.title = titulo
        ws['A1'], ws['A2'] = DC.TITULO, cabecera
        assert ws['B4'].value and 'TRÁMITE' in ws['B4'].value.upper(), ws['B4'].value
        modelo = [copy.copy(ws.cell(5, c)) for c in range(1, 6)]
        alto = ws.row_dimensions[5].height
        pie = [copy.copy(ws.cell(r, c)) for r in range(5, ws.max_row + 1) for c in range(1, 6)
               if isinstance(ws.cell(r, c).value, str) and 'completad' in ws.cell(r, c).value.lower()]
        assert pie, ws.title
        fila_pie = pie[0].row
        estilo_pie = [copy.copy(ws.cell(fila_pie, c)) for c in range(1, 6)]
        # vaciar la zona de tareas y el pie (valores, DV y formato condicional los rehace el motor)
        for r in range(5, ws.max_row + 1):
            for c in range(1, 6):
                ws.cell(r, c).value = None
            ws.row_dimensions[r].height = None
        ws.data_validations.dataValidation = []
        ws.conditional_formatting = type(ws.conditional_formatting)()
        for i, (tarea, resp, plazo, nota) in enumerate(tareas):
            r = 5 + i
            for c, v in enumerate(('☐', tarea, resp, plazo, nota), start=1):
                cel = ws.cell(r, c)
                _copiar_estilo(modelo[c - 1], cel)
                cel.value = v
            ws.row_dimensions[r].height = alto
        r = 5 + len(tareas) + 1
        for c in range(1, 6):
            _copiar_estilo(estilo_pie[c - 1], ws.cell(r, c))
        ws.cell(r, 2).value = 'Tareas completadas:'
        ws.cell(r, 3).value = '=COUNTIF(A5:A%d,"✓")' % (4 + len(tareas))
        ws.cell(r, 4).value = 'de'
        ws.cell(r, 5).value = '=COUNTIF(B5:B%d,"?*")-COUNTIF(A5:A%d,"N/A")' % (4 + len(tareas), 4 + len(tareas))
    # `motor.asegurar_instrucciones` no reescribe la hoja si ya existe: la de la panadería nombraba su producto
    ins = wb['Instrucciones']
    assert 'plan-negocio-panaderia' in str(ins['A3'].value), ins['A3'].value
    ins['A3'] = 'Este fichero forma parte del producto «%s» de AI Chef Pro.' % PID
    _docprops(wb)
    wb.save(destino)
    return destino


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--salida', required=True)
    a = ap.parse_args()
    if os.path.isdir(a.salida):
        shutil.rmtree(a.salida)
    os.makedirs(a.salida)
    print('plan      →', plan(a.salida))
    print('checklist →', checklist(a.salida), '(%d tareas: %s)' % (
        sum(len(t) for _, _, t in DC.FASES), '/'.join(str(len(t)) for _, _, t in DC.FASES)))


if __name__ == '__main__':
    main()
