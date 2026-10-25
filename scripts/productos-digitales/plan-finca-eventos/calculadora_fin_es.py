#!/usr/bin/env python3
"""
calculadora_fin_es.py — Plan de Negocio: Finca de Eventos y Bodas ES 1.0 (SPEC D10): `calculadora-presupuesto-boda-evento-finca.xlsx`,
del calendario y la tipología del evento al presupuesto del ESPACIO. Fork de `parrillero-2_2/calculadora_par_es.py`:
  · canal Particular / Empresa (desplegable) y tipología (boda, empresa, celebración privada, celebración pequeña);
  · la FECHA fija la temporada (alta / media / baja, tabla de meses) y el día (sábado / viernes o domingo / entre semana,
    tabla de días de la semana); el ajuste de cada una es un suplemento o un descuento en una tabla VISIBLE y editable;
  · mínimo de invitados por temporada (invitados garantizados = el mayor de los dos);
  · servicios extra de la finca (ceremonia, horas extra, montaje premium, alojamiento): precios de EJEMPLO en celdas verdes;
  · IVA al 21 % del espacio y los servicios (la finca NO factura el menú: lo factura el catering exclusivo, aparte);
  · canon del catering estimado (INFORMATIVO: % editable × menú por invitado × invitados; no lo paga el cliente a la finca);
  · señal y calendario de pagos (reserva, segundo pago, liquidación) con sus fechas límite;
  · cancelación por tramos en MESES SIN huecos, proporcionados (TRLGDCU art. 85.6) y recíprocos (art. 87.2);
  · margen del evento para la finca CALCULADO (coste directo del evento editable); 3 ejemplos por las mismas fórmulas;
    autotest pycel ± 0,01 €.
Hojas `Calculadora` · `Ejemplos` · `Instrucciones` (firma `calc-beta` de `motor.detectar`). A4, protegida sin contraseña,
verdes desbloqueadas. Corre en el VPS (openpyxl + pycel).

    /root/venv-guias/bin/python calculadora_fin_es.py --salida <dir>      # genera + inject_cache + autotest
    /root/venv-guias/bin/python calculadora_fin_es.py --salida <dir> --autotest

Sesión Claude Code, 25-oct-2026.
"""
import argparse
import datetime
import math
import os
import re
import sys
from collections import OrderedDict

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill, Protection
from openpyxl.worksheet.datavalidation import DataValidation

AQUI = os.path.dirname(os.path.abspath(__file__))
PD = os.path.dirname(AQUI)
sys.path.insert(0, AQUI)
import fin_version                                                                 # noqa: E402,F401  (parchea motor: 1.0 · noviembre 2026)
import motor                                                                       # noqa: E402

PID = 'plan-negocio-finca-eventos-bodas'
FICHERO = 'calculadora-presupuesto-boda-evento-finca.xlsx'
TITULO = 'Calculadora de presupuesto del espacio para bodas y eventos en finca'
H1, H2, H3 = 'Calculadora', 'Ejemplos', 'Instrucciones'
VERDE = PatternFill('solid', fgColor='E8F5E9')
GRIS = PatternFill('solid', fgColor='F2F2F2')
FUENTE = 'Calibri'
CANALES = ['Particular', 'Empresa']
EUR, ENT, PCT, DEC, FECHA = '#,##0.00 €', '#,##0', '0%', '#,##0.0', 'dd/mm/yyyy'
EPOCA = datetime.date(1899, 12, 30)


def serial(d):
    """Fecha → número de serie de Excel (así se guardan las fechas: pycel y Excel leen lo mismo)."""
    return (d - EPOCA).days


# ---------------------------------------------------------------------------------------------------------------------
# Tablas de datos (F1-research §2 y §3). Son la FUENTE ÚNICA: la calculadora las escribe en celdas verdes.
# Todas las cifras son ESTIMACIONES DE EJEMPLO del kit: cámbialas por las de tu finca.
# ---------------------------------------------------------------------------------------------------------------------
#: tipología → (tarifa del espacio sábado de temporada media, sin IVA · menú del catering por invitado, sin IVA · coste directo del evento)
TIPOS = OrderedDict([
    ('Boda', (3600, 110.0, 650)), ('Empresa', (2300, 42.0, 350)),
    ('Celebración privada', (1600, 50.0, 300)), ('Celebración pequeña', (1050, 33.0, 150))])
#: temporada → (ajuste sobre la tarifa, mínimo de invitados garantizados)
TEMPORADAS = OrderedDict([('Alta', (0.20, 80)), ('Media', (0.0, 60)), ('Baja', (-0.20, 40))])
#: tipo de día → ajuste sobre la tarifa
DIAS = OrderedDict([('Sábado', 0.0), ('Viernes o domingo', -0.20), ('Entre semana', -0.30)])
#: mes → temporada (INE: mayo-octubre concentra el 63,7 % de los matrimonios; septiembre y junio son los picos)
MESES = [('Enero', 'Baja'), ('Febrero', 'Baja'), ('Marzo', 'Baja'), ('Abril', 'Media'), ('Mayo', 'Alta'), ('Junio', 'Alta'),
         ('Julio', 'Media'), ('Agosto', 'Media'), ('Septiembre', 'Alta'), ('Octubre', 'Alta'), ('Noviembre', 'Baja'),
         ('Diciembre', 'Baja')]
#: día de la semana (domingo = 1, como WEEKDAY de Excel) → tipo de día
SEMANA = [('Domingo', 'Viernes o domingo'), ('Lunes', 'Entre semana'), ('Martes', 'Entre semana'), ('Miércoles', 'Entre semana'),
          ('Jueves', 'Entre semana'), ('Viernes', 'Viernes o domingo'), ('Sábado', 'Sábado')]
#: servicio extra de la finca → precio de EJEMPLO sin IVA
EXTRAS = OrderedDict([
    ('cer', ('Ceremonia civil o simbólica en la finca (por ceremonia)', 450.0)),
    ('hex', ('Hora extra de uso del espacio (por hora)', 300.0)),
    ('mon', ('Montaje premium: decoración, mobiliario y luces (por montaje)', 600.0)),
    ('alo', ('Alojamiento, si la finca lo tiene (por habitación y noche)', 150.0))])
#: política por canal: Particular · Empresa
CANAL_TAB = OrderedDict([
    ('pc_1', ('Cobro 1: reserva de la fecha (% del total)', (0.25, 0.30),
              'Al firmar. Es la señal: tiene que ser proporcionada al coste de bloquear la fecha (art. 85.6 del TRLGDCU).')),
    ('pc_2', ('Cobro 2: segundo pago (% del total)', (0.35, 0.30), 'Pago intermedio, en la fecha de la fila siguiente.')),
    ('pc_3', ('Cobro 3: liquidación antes del evento (% del total)', (0.40, 0.40),
              'El resto, antes del evento. Emite factura de cada pago, también de la reserva.')),
    ('d_2', ('Segundo pago: días antes del evento', (180, 60),
             'Particular, unos seis meses antes; empresa, dos meses. Si firman con menos antelación, se cobran juntos.')),
    ('d_3', ('Liquidación: días antes del evento', (30, 15), 'Particular, 30 días antes; empresa, 15.')),
    ('cn1', ('Cancelación tramo 1: lo que retienes (% del total)', (0.10, 0.05),
             'Tramo más lejano del evento. Nunca más de lo cobrado en la reserva.')),
    ('cn2', ('Cancelación tramo 2: lo que retienes (% del total)', (0.25, 0.15), 'Sube con la cercanía del evento.')),
    ('cn3', ('Cancelación tramo 3: lo que retienes (% del total)', (0.50, 0.35), 'Cada tramo repone menos opciones de volver a vender la fecha.')),
    ('cn4', ('Cancelación tramo 4: lo que retienes (% del total)', (0.75, 0.60), 'Últimos meses.')),
    ('cn5', ('Cancelación tramo 5: lo que retienes (% del total)', (1.0, 1.0), 'Último mes o no presentación: el 100 %.')),
])
#: tramos de cancelación: meses de antelación a partir de los cuales aplica cada uno (de más lejano a más cercano, sin huecos)
TRAMOS = [12, 6, 3, 1, 0]

ENTRADAS = [
    ('canal', 'Cliente', '@', 'Particular',
     'Particular o Empresa: cambian el calendario de cobros y los tramos de cancelación (tabla del final). Una empresa '
     'recibe el presupuesto «más IVA»; un particular, con el IVA incluido.'),
    ('tipo', 'Tipología del evento', '@', 'Boda',
     'Boda, empresa (jornada, comida, Navidad), celebración privada (comunión, bautizo, cumpleaños) o celebración pequeña (preboda, brunch, boda íntima). Fija la tarifa del '
     'espacio, el menú del catering y el coste directo con las tablas del final.'),
    ('fecha', 'Fecha del evento', FECHA, serial(datetime.date(2027, 6, 19)),
     'Escribe la fecha completa (dd/mm/aaaa). Fija la temporada y el tipo de día con las tablas de meses y de días de la semana.'),
    ('guests', 'Invitados previstos', ENT, 110, 'Número entero de invitados. Se garantiza el mayor entre este número y el mínimo de la temporada.'),
    ('n_cer', 'Ceremonias en la finca (nº)', ENT, 1, 'Civil o simbólica, a precio de la tabla de servicios extra.'),
    ('n_hex', 'Horas extra de uso del espacio (nº)', ENT, 2, 'Las horas que pasan del horario base de tu contrato.'),
    ('n_mon', 'Montajes premium (nº)', ENT, 1, 'Decoración, mobiliario y luces por encima del montaje básico.'),
    ('n_alo', 'Alojamiento: habitaciones-noche (nº)', ENT, 0, '0 si la finca no tiene habitaciones. Ejemplo: 6 habitaciones × 2 noches = 12.'),
    ('h_aux', 'Horas de personal de la finca en el evento (coordinación y auxiliares)', DEC, 16,
     'Estimación del kit: pon las horas reales de tu equipo (no incluye cocina ni sala: las pone el catering).'),
    ('wage', 'Coste bruto por hora del personal', EUR, 13.0, 'Estimación del kit. Copia tu convenio provincial de hostelería.'),
    ('ss', 'Seguridad Social a cargo de la empresa', PCT, 0.33,
     '30,65 % de contingencias, desempleo, FOGASA, formación y MEI, más AT/EP (fijos discontinuos, art. 16 ET).'),
    ('cost_ext', 'Coste directo de los servicios extra (% sobre su precio)', PCT, 0.40,
     'Estimación del kit: lo que te cuesta prestar la ceremonia, el montaje, las horas extra y el alojamiento.'),
    ('redondeo', 'Redondeo del precio del espacio (múltiplo, en euros)', ENT, 50, 'El precio del espacio se redondea a este múltiplo, al más cercano.'),
    ('canon', 'Canon del catering exclusivo (% de lo que factura el catering)', PCT, 0.12,
     'SUPUESTO editable, no un dato: negócialo con tu catering. Informativo: lo paga el catering a la finca, no el cliente.'),
    ('iva', 'IVA de la finca (espacio, servicios y canon)', PCT, 0.21,
     'El espacio y los servicios tienen precio propio: la finca los factura al 21 %. El canon también lleva IVA.'),
    ('iva_menu', 'IVA del menú del catering (solo orientativo)', PCT, 0.10,
     'El menú lo factura el catering, no la finca (comida y bebida para consumir en el acto, art. 91.Uno.2.2.º de la Ley 37/1992).'),
]
K = OrderedDict((k, 5 + i) for i, (k, *_r) in enumerate(ENTRADAS))
F_CALC = 5 + len(ENTRADAS) + 1

# nombre, rótulo, formato, expresión, nota. {x}: celda de la entrada/salida; {t_xxx}: búsquedas en las tablas
SALIDAS = [
    ('temporada', 'Temporada de la fecha (de la tabla de meses)', 'General', '{t_temporada}', 'Busca el mes de la fecha en la tabla del final.'),
    ('dia', 'Tipo de día (de la tabla de días de la semana)', 'General', '{t_dia}', 'Busca el día de la semana de la fecha en la tabla del final.'),
    ('aj_t', 'Ajuste de temporada (suplemento + / descuento −)', PCT, '{t_aj_t}', 'De la tabla de temporadas.'),
    ('aj_d', 'Ajuste de día (suplemento + / descuento −)', PCT, '{t_aj_d}', 'De la tabla de días.'),
    ('tarifa', 'Tarifa del espacio de la tipología (sábado, temporada media)', EUR, '{t_tarifa}', 'De la tabla de tipologías.'),
    ('espacio', 'ALQUILER DEL ESPACIO SIN IVA', EUR, 'ROUND({tarifa}*(1+{aj_t})*(1+{aj_d})/{redondeo},0)*{redondeo}',
     'Tarifa × (1 + ajuste de temporada) × (1 + ajuste de día), redondeada al múltiplo elegido.'),
    ('min_inv', 'Mínimo de invitados de la temporada', ENT, '{t_min}', 'De la tabla de temporadas.'),
    ('garant', 'Invitados garantizados', ENT, 'MAX({guests},{min_inv})', 'El mayor entre los invitados previstos y el mínimo de la temporada.'),
    ('aviso', 'Aviso del mínimo', 'General', 'IF({guests}<{min_inv},"Por debajo del mínimo: se garantiza el mínimo de la temporada","Cumple el mínimo")', ''),
    ('e_cer', 'Ceremonias', EUR, '{n_cer}*{x_cer}', 'Nº × precio de la tabla de servicios extra.'),
    ('e_hex', 'Horas extra', EUR, '{n_hex}*{x_hex}', ''),
    ('e_mon', 'Montajes premium', EUR, '{n_mon}*{x_mon}', ''),
    ('e_alo', 'Alojamiento', EUR, '{n_alo}*{x_alo}', ''),
    ('extras', 'Servicios extra sin IVA', EUR, '{e_cer}+{e_hex}+{e_mon}+{e_alo}', ''),
    ('base', 'Base imponible (espacio + servicios)', EUR, '{espacio}+{extras}', ''),
    ('iva_amt', 'IVA de la finca', EUR, '{base}*{iva}', ''),
    ('total', 'TOTAL DE LA FINCA (con IVA)', EUR, '{base}+{iva_amt}', 'No incluye el menú ni la bebida: los factura el catering exclusivo, aparte.'),
    ('quoted', 'Importe que presupuestas (Particular: con IVA · Empresa: más IVA)', EUR, 'IF({canal}="Particular",{total},{base})',
     'A una empresa se le presupuesta sin IVA («+ IVA»); a un particular, con los impuestos incluidos.'),
    ('eur_inv', 'Precio del espacio y servicios por invitado garantizado (sin IVA)', EUR, '{base}/{garant}', 'Para comparar con otras fincas.'),
    ('cb1', 'Importe del cobro 1: reserva', EUR, '{total}*{t_pc_1}', 'Al firmar el contrato.'),
    ('cb2', 'Importe del cobro 2: segundo pago', EUR, '{total}*{t_pc_2}', ''),
    ('f_2', 'Fecha límite del segundo pago', FECHA, '{fecha}-{t_d_2}', 'Fecha del evento menos los días de la tabla.'),
    ('cb3', 'Importe del cobro 3: liquidación', EUR, '{total}*{t_pc_3}', ''),
    ('f_3', 'Fecha límite de la liquidación', FECHA, '{fecha}-{t_d_3}', 'Fecha del evento menos los días de la tabla.'),
    ('menu_tot', 'Importe del menú del catering sin IVA (estimado, invitados garantizados)', EUR, '{t_menu}*{garant}',
     'INFORMATIVO. Menú por invitado de la tabla de tipologías × invitados garantizados. No es ingreso de la finca.'),
    ('canon_est', 'CANON ESTIMADO DEL CATERING sin IVA (informativo)', EUR, '{menu_tot}*{canon}',
     'Lo paga el catering exclusivo a la finca, no el cliente. Es un supuesto: negócialo.'),
    ('canon_iva', 'Canon estimado con IVA', EUR, '{canon_est}*(1+{iva})', 'La finca factura el canon al catering con IVA.'),
    ('menu_cli', 'Menú del catering con IVA (orientativo, lo cobra el catering)', EUR, '{menu_tot}*(1+{iva_menu})', 'Para que el cliente vea el coste completo.'),
    ('presup', 'Importe orientativo completo para el cliente (finca con IVA + menú)', EUR, '{total}+{menu_cli}', 'Suma lo que paga a la finca y lo que paga al catering.'),
    ('c_fijo', 'Coste directo del evento: limpieza, aseos, gasoil, vigilancia y música', EUR, '{t_cfijo}',
     'De la tabla de tipologías (estimación del kit, sin tarifas públicas).'),
    ('c_aux', 'Coste de personal de la finca (salario + Seguridad Social)', EUR, '{h_aux}*{wage}*(1+{ss})', ''),
    ('c_ext', 'Coste directo de los servicios extra', EUR, '{extras}*{cost_ext}', ''),
    ('c_total', 'Coste directo total del evento', EUR, '{c_fijo}+{c_aux}+{c_ext}', ''),
    ('ingresos', 'Ingresos de la finca por el evento (sin IVA, con canon)', EUR, '{base}+{canon_est}', 'Espacio + servicios + canon estimado.'),
    ('profit', 'Margen del evento para la finca (euros)', EUR, '{ingresos}-{c_total}',
     'Contribución: antes de la renta, el personal fijo, los suministros, los seguros y la estructura del plan financiero.'),
    ('margin_calc', 'Margen del evento (calculado)', '0.0%', 'IF({ingresos}=0,0,{profit}/{ingresos})', 'Sale de los costes, nunca se teclea.'),
    ('profit_sc', 'Margen del evento sin el canon (euros)', EUR, '{base}-{c_total}', 'Cuánto dejaría el evento si el catering no pagara canon.'),
    ('can1', 'Importe retenido si cancela: tramo 1', EUR, '{total}*{t_cn1}', 'Tramo más lejano al evento (ver tabla de tramos). Si cancela la finca: devuelve lo cobrado y paga lo mismo.'),
    ('can2', 'Importe retenido si cancela: tramo 2', EUR, '{total}*{t_cn2}', ''),
    ('can3', 'Importe retenido si cancela: tramo 3', EUR, '{total}*{t_cn3}', ''),
    ('can4', 'Importe retenido si cancela: tramo 4', EUR, '{total}*{t_cn4}', ''),
    ('can5', 'Importe retenido si cancela: tramo 5', EUR, '{total}*{t_cn5}', 'Último tramo, hasta el día del evento.'),
]
S = OrderedDict((k, F_CALC + 1 + i) for i, (k, *_r) in enumerate(SALIDAS))

# --- bloques del pie (tablas) ---------------------------------------------------------------------------------------
F_TIPOS = F_CALC + 1 + len(SALIDAS) + 1
F_TEMP = F_TIPOS + 1 + len(TIPOS) + 1
F_DIA = F_TEMP + 1 + len(TEMPORADAS) + 1
F_MES = F_DIA + 1 + len(DIAS) + 1
F_SEM = F_MES + 1 + len(MESES) + 1
F_EXT = F_SEM + 1 + len(SEMANA) + 1
X = OrderedDict((k, F_EXT + 1 + i) for i, k in enumerate(EXTRAS))
F_CAN = F_EXT + 1 + len(EXTRAS) + 1
CT = OrderedDict((k, F_CAN + 1 + i) for i, k in enumerate(CANAL_TAB))
F_TRA = F_CAN + 1 + len(CANAL_TAB) + 1

EJEMPLOS = OrderedDict([
    ('Boda de sábado en temporada alta', dict(
        canal='Particular', tipo='Boda', fecha=serial(datetime.date(2027, 6, 19)), guests=110, n_cer=1, n_hex=2, n_mon=1,
        n_alo=0, h_aux=16, wage=13.0, ss=0.33, cost_ext=0.40, redondeo=50, canon=0.12, iva=0.21, iva_menu=0.10)),
    ('Boda de viernes en temporada media', dict(
        canal='Particular', tipo='Boda', fecha=serial(datetime.date(2027, 7, 16)), guests=90, n_cer=1, n_hex=0, n_mon=0,
        n_alo=12, h_aux=14, wage=13.0, ss=0.33, cost_ext=0.40, redondeo=50, canon=0.12, iva=0.21, iva_menu=0.10)),
    ('Jornada de empresa entre semana', dict(
        canal='Empresa', tipo='Empresa', fecha=serial(datetime.date(2027, 4, 21)), guests=50, n_cer=0, n_hex=1,
        n_mon=0, n_alo=0, h_aux=10, wage=13.0, ss=0.33, cost_ext=0.40, redondeo=50, canon=0.12, iva=0.21, iva_menu=0.10)),
])
EJ_NOTAS = OrderedDict([
    ('Boda de sábado en temporada alta', 'Particular, 110 invitados, un sábado de junio: temporada alta (+20 %) y día de referencia '
     '(sábado, sin ajuste). Lleva ceremonia en la finca, dos horas extra y montaje premium. El precio con IVA es el que ve el '
     'particular. El canon es lo que el catering exclusivo pagaría a la finca con un 12 % (supuesto): no lo paga el novio. '
     'Cobro en tres pagos: reserva, segundo pago seis meses antes y liquidación 30 días antes.'),
    ('Boda de viernes en temporada media', 'Particular, 90 invitados, un viernes de julio: temporada media (sin ajuste) y viernes '
     '(−20 %). Lleva ceremonia y 12 habitaciones-noche de alojamiento (6 habitaciones, 2 noches). Mismo calendario de cobros y '
     'mismos tramos que la anterior: cambia solo la tarifa del espacio.'),
    ('Jornada de empresa entre semana', 'Empresa, 50 invitados, un miércoles de abril: temporada media y día entre semana (−30 %). '
     'Los 50 invitados quedan por debajo del mínimo de la temporada (60): se garantizan 60 y el canon se calcula sobre ellos. Se presupuesta '
     '«más IVA» y se cobra con otro calendario (30 % al reservar, 30 % a dos meses y 40 % a 15 días) y otros tramos de cancelación.'),
])


# ---------------------------------------------------------------------------------------------------------------------
# Cálculo en Python (espejo de las fórmulas) — la fuente del autotest
# ---------------------------------------------------------------------------------------------------------------------
def _r(x):
    """ROUND de Excel (la mitad se aleja de cero)."""
    return math.floor(x + 0.5) if x >= 0 else -math.floor(-x + 0.5)


def _fecha(n):
    return EPOCA + datetime.timedelta(days=int(n))


def calcular(e, tablas=None):
    r = OrderedDict()
    ic = CANALES.index(e['canal'])
    t = lambda k: CANAL_TAB[k][1][ic]
    d = _fecha(e['fecha'])
    r['temporada'] = MESES[d.month - 1][1]
    r['dia'] = SEMANA[d.isoweekday() % 7][1]
    r['aj_t'] = TEMPORADAS[r['temporada']][0]
    r['aj_d'] = DIAS[r['dia']]
    tarifa, menu, cfijo = TIPOS[e['tipo']]
    r['tarifa'] = tarifa
    r['espacio'] = _r(tarifa * (1 + r['aj_t']) * (1 + r['aj_d']) / e['redondeo']) * e['redondeo']
    r['min_inv'] = TEMPORADAS[r['temporada']][1]
    r['garant'] = max(e['guests'], r['min_inv'])
    r['aviso'] = ('Por debajo del mínimo: se garantiza el mínimo de la temporada' if e['guests'] < r['min_inv']
                  else 'Cumple el mínimo')
    r['e_cer'] = e['n_cer'] * EXTRAS['cer'][1]
    r['e_hex'] = e['n_hex'] * EXTRAS['hex'][1]
    r['e_mon'] = e['n_mon'] * EXTRAS['mon'][1]
    r['e_alo'] = e['n_alo'] * EXTRAS['alo'][1]
    r['extras'] = r['e_cer'] + r['e_hex'] + r['e_mon'] + r['e_alo']
    r['base'] = r['espacio'] + r['extras']
    r['iva_amt'] = r['base'] * e['iva']
    r['total'] = r['base'] + r['iva_amt']
    r['quoted'] = r['total'] if e['canal'] == 'Particular' else r['base']
    r['eur_inv'] = r['base'] / r['garant']
    r['cb1'] = r['total'] * t('pc_1')
    r['cb2'] = r['total'] * t('pc_2')
    r['f_2'] = e['fecha'] - t('d_2')
    r['cb3'] = r['total'] * t('pc_3')
    r['f_3'] = e['fecha'] - t('d_3')
    r['menu_tot'] = menu * r['garant']
    r['canon_est'] = r['menu_tot'] * e['canon']
    r['canon_iva'] = r['canon_est'] * (1 + e['iva'])
    r['menu_cli'] = r['menu_tot'] * (1 + e['iva_menu'])
    r['presup'] = r['total'] + r['menu_cli']
    r['c_fijo'] = cfijo
    r['c_aux'] = e['h_aux'] * e['wage'] * (1 + e['ss'])
    r['c_ext'] = r['extras'] * e['cost_ext']
    r['c_total'] = r['c_fijo'] + r['c_aux'] + r['c_ext']
    r['ingresos'] = r['base'] + r['canon_est']
    r['profit'] = r['ingresos'] - r['c_total']
    r['margin_calc'] = 0 if r['ingresos'] == 0 else r['profit'] / r['ingresos']
    r['profit_sc'] = r['base'] - r['c_total']
    for i in (1, 2, 3, 4, 5):
        r['can%d' % i] = r['total'] * t('cn%d' % i)
    return r


# ---------------------------------------------------------------------------------------------------------------------
# Fórmulas
# ---------------------------------------------------------------------------------------------------------------------
def _rng(col, f1, f2):
    return "'%s'!$%s$%d:$%s$%d" % (H1, col, f1, col, f2)


def _busca(col, f0, n, clave):
    return 'INDEX(%s,MATCH(%s,%s,0))' % (_rng(col, f0 + 1, f0 + n), clave, _rng('A', f0 + 1, f0 + n))


def refs(col):
    """Referencias a celdas de la columna `col` (B en Calculadora; B/C/D en Ejemplos) + búsquedas en las tablas."""
    ref = dict((k, '$%s$%d' % (col, f)) for k, f in K.items())
    ref.update((k, '$%s$%d' % (col, f)) for k, f in S.items())
    ref['t_tarifa'] = _busca('B', F_TIPOS, len(TIPOS), ref['tipo'])
    ref['t_menu'] = _busca('C', F_TIPOS, len(TIPOS), ref['tipo'])
    ref['t_cfijo'] = _busca('D', F_TIPOS, len(TIPOS), ref['tipo'])
    ref['t_aj_t'] = _busca('B', F_TEMP, len(TEMPORADAS), ref['temporada'])
    ref['t_min'] = _busca('C', F_TEMP, len(TEMPORADAS), ref['temporada'])
    ref['t_aj_d'] = _busca('B', F_DIA, len(DIAS), ref['dia'])
    ref['t_temporada'] = 'INDEX(%s,MONTH(%s))' % (_rng('B', F_MES + 1, F_MES + len(MESES)), ref['fecha'])
    ref['t_dia'] = 'INDEX(%s,WEEKDAY(%s))' % (_rng('B', F_SEM + 1, F_SEM + len(SEMANA)), ref['fecha'])
    for k, f in X.items():
        ref['x_' + k] = "'%s'!$B$%d" % (H1, f)
    cab = "'%s'!$B$%d:$C$%d" % (H1, F_CAN, F_CAN)
    for k, f in CT.items():
        ref['t_' + k] = "INDEX('%s'!$B$%d:$C$%d,MATCH(%s,%s,0))" % (H1, f, f, ref['canal'], cab)
    return ref


def a_excel(expr, ref):
    return re.sub(r'\{(\w+)\}', lambda m: ref[m.group(1)], expr)


def _f(ws, coord, v, b=False, sz=11, i=False, fmt=None, wrap=False, fill=None, locked=True):
    c = ws[coord]
    c.value = v
    c.font = Font(name=FUENTE, size=sz, bold=b, italic=i)
    if fmt:
        c.number_format = fmt
    c.alignment = Alignment(wrap_text=wrap, vertical='top')
    if fill:
        c.fill = fill
    c.protection = Protection(locked=locked)
    return c


def _nota(ws, fila, texto, desde='C'):
    """Nota combinada en C:F (ancho útil ≈ 110 caracteres)."""
    if not texto:
        return
    _f(ws, '%s%d' % (desde, fila), texto, wrap=True)
    ws.merge_cells('%s%d:F%d' % (desde, fila, fila))
    ws.row_dimensions[fila].height = max(16, 15 * math.ceil(len(texto) / (105.0 if desde == 'C' else 80.0)))


def _imprimir(ws, horizontal=False):
    ws.page_setup.paperSize = 9                                                   # A4
    ws.page_setup.orientation = 'landscape' if horizontal else 'portrait'
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.oddFooter.center.text = 'aichef.pro/%s · página &P de &N' % PID
    ws.print_options.horizontalCentered = True
    ws.protection.sheet = True                                                    # sin contraseña
    ws.sheet_view.showGridLines = False


def _dv(ws, tipo, celdas, f1=None, f2=None, op=None, titulo='Valor no válido', msg=''):
    dv = DataValidation(type=tipo, formula1=f1, formula2=f2, operator=op, allow_blank=False, showErrorMessage=True,
                        errorTitle=titulo, error=msg)
    ws.add_data_validation(dv)
    for c in celdas:
        dv.add(c)


def _cab(ws, f, titulo, cols):
    _f(ws, 'A%d' % f, titulo, b=True, fill=GRIS)
    for col, t in cols:
        _f(ws, '%s%d' % (col, f), t, b=True, fill=GRIS, wrap=True)
    ws.row_dimensions[f].height = 32


def hoja_calculadora(wb):
    ws = wb.active
    ws.title = H1
    for col, w in (('A', 66), ('B', 22), ('C', 17), ('D', 17), ('E', 17), ('F', 56)):
        ws.column_dimensions[col].width = w
    _f(ws, 'A1', 'CALCULADORA DE PRESUPUESTO DEL ESPACIO PARA BODAS Y EVENTOS EN FINCA', b=True, sz=14)
    _f(ws, 'A2', 'Escribe sólo en las celdas VERDES: todo lo demás se calcula. Importes sin IVA salvo donde la fila '
                 'diga lo contrario. Los valores por defecto son estimaciones de ejemplo del kit: cámbialos por los de tu finca.',
       wrap=True)
    ws.merge_cells('A2:F2')
    ws.row_dimensions[2].height = 32
    for col, t in zip('ABC', ('DATOS DEL EVENTO', 'Valor', 'Notas')):
        _f(ws, '%s4' % col, t, b=True, fill=GRIS)
    for k, rot, fmt, val, nota in ENTRADAS:
        f = K[k]
        _f(ws, 'A%d' % f, rot)
        _f(ws, 'B%d' % f, val, fmt=fmt if fmt != '@' else None, fill=VERDE, locked=False)
        _nota(ws, f, nota)
    for col, t in zip('ABC', ('CÁLCULO', 'Valor', 'Notas')):
        _f(ws, '%s%d' % (col, F_CALC), t, b=True, fill=GRIS)
    ref = refs('B')
    for k, rot, fmt, expr, nota in SALIDAS:
        f = S[k]
        neg = k in ('espacio', 'total', 'quoted', 'canon_est', 'margin_calc', 'presup')
        _f(ws, 'A%d' % f, rot, b=neg)
        _f(ws, 'B%d' % f, '=' + a_excel(expr, ref), fmt=fmt, b=neg)
        _nota(ws, f, nota)
    # tipologías
    _cab(ws, F_TIPOS, 'TIPOLOGÍA → TARIFA, MENÚ DEL CATERING Y COSTE DIRECTO (sin IVA)',
         [('B', 'Tarifa del espacio (sábado, temporada media)'), ('C', 'Menú del catering por invitado'),
          ('D', 'Coste directo del evento')])
    for i, (nombre, (tarifa, menu, cfijo)) in enumerate(TIPOS.items(), 1):
        _f(ws, 'A%d' % (F_TIPOS + i), nombre)
        _f(ws, 'B%d' % (F_TIPOS + i), tarifa, fmt=EUR, fill=VERDE, locked=False)
        _f(ws, 'C%d' % (F_TIPOS + i), menu, fmt=EUR, fill=VERDE, locked=False)
        _f(ws, 'D%d' % (F_TIPOS + i), cfijo, fmt=EUR, fill=VERDE, locked=False)
    _f(ws, 'E%d' % (F_TIPOS + 1), 'Cifras de EJEMPLO del kit (precios de escaparate y estimaciones, no tarifas de mercado). '
                                  'Coste directo: limpieza, aseos, gasoil del grupo, vigilancia y música (SGAE/AGEDI).', i=True, wrap=True)
    ws.merge_cells('E%d:F%d' % (F_TIPOS + 1, F_TIPOS + 4))
    # temporadas
    _cab(ws, F_TEMP, 'TEMPORADA → AJUSTE DE TARIFA Y MÍNIMO DE INVITADOS',
         [('B', 'Ajuste (+ suplemento / − descuento)'), ('C', 'Mínimo de invitados')])
    for i, (nombre, (aj, mn)) in enumerate(TEMPORADAS.items(), 1):
        _f(ws, 'A%d' % (F_TEMP + i), nombre)
        _f(ws, 'B%d' % (F_TEMP + i), aj, fmt=PCT, fill=VERDE, locked=False)
        _f(ws, 'C%d' % (F_TEMP + i), mn, fmt=ENT, fill=VERDE, locked=False)
    _f(ws, 'D%d' % (F_TEMP + 1), 'Ejemplo del kit. La temporada alta cobra suplemento y la baja descuenta; la tarifa de las '
                                 'tipologías es la de un sábado de temporada media.', i=True, wrap=True)
    ws.merge_cells('D%d:F%d' % (F_TEMP + 1, F_TEMP + 3))
    # días
    _cab(ws, F_DIA, 'TIPO DE DÍA → AJUSTE DE TARIFA', [('B', 'Ajuste (+ suplemento / − descuento)')])
    for i, (nombre, aj) in enumerate(DIAS.items(), 1):
        _f(ws, 'A%d' % (F_DIA + i), nombre)
        _f(ws, 'B%d' % (F_DIA + i), aj, fmt=PCT, fill=VERDE, locked=False)
    _f(ws, 'C%d' % (F_DIA + 1), 'Ejemplo del kit: las fichas de fincas publican descuentos de hasta el 30 % en viernes, domingo y de '
                                'noviembre a marzo. Los ajustes se multiplican entre sí.', i=True, wrap=True)
    ws.merge_cells('C%d:F%d' % (F_DIA + 1, F_DIA + 3))
    # meses
    _cab(ws, F_MES, 'MES → TEMPORADA', [('B', 'Temporada')])
    for i, (mes, temp) in enumerate(MESES, 1):
        _f(ws, 'A%d' % (F_MES + i), mes)
        _f(ws, 'B%d' % (F_MES + i), temp, fill=VERDE, locked=False)
    _f(ws, 'C%d' % (F_MES + 1), 'Mayo-octubre concentra el 63,7 % de los matrimonios y septiembre y junio son los picos '
                                '(INE, 2024). Ajusta los meses a tu zona.', i=True, wrap=True)
    ws.merge_cells('C%d:F%d' % (F_MES + 1, F_MES + 3))
    # semana
    _cab(ws, F_SEM, 'DÍA DE LA SEMANA → TIPO DE DÍA', [('B', 'Tipo de día')])
    for i, (dia, tipo) in enumerate(SEMANA, 1):
        _f(ws, 'A%d' % (F_SEM + i), dia)
        _f(ws, 'B%d' % (F_SEM + i), tipo, fill=VERDE, locked=False)
    # extras
    _cab(ws, F_EXT, 'SERVICIOS EXTRA DE LA FINCA (precio sin IVA)', [('B', 'Precio')])
    for k, (rot, p) in EXTRAS.items():
        _f(ws, 'A%d' % X[k], rot)
        _f(ws, 'B%d' % X[k], p, fmt=EUR, fill=VERDE, locked=False)
    _f(ws, 'C%d' % (F_EXT + 1), 'Precios de EJEMPLO del kit: pon los de tu finca. El alojamiento sólo si la finca tiene habitaciones.',
       i=True, wrap=True)
    ws.merge_cells('C%d:F%d' % (F_EXT + 1, F_EXT + 2))
    # por cliente
    _cab(ws, F_CAN, 'POLÍTICA POR CLIENTE', [('B', CANALES[0]), ('C', CANALES[1])])
    for k, (rot, vals, nota) in CANAL_TAB.items():
        f = CT[k]
        _f(ws, 'A%d' % f, rot)
        for col, v in zip('BC', vals):
            _f(ws, '%s%d' % (col, f), v, fmt=ENT if k.startswith('d_') else PCT, fill=VERDE, locked=False)
        _f(ws, 'D%d' % f, nota, wrap=True)
        ws.merge_cells('D%d:F%d' % (f, f))
        ws.row_dimensions[f].height = max(16, 15 * math.ceil(len(nota) / 78.0))
    # tramos de cancelación
    _cab(ws, F_TRA, 'TRAMOS DE CANCELACIÓN (meses de antelación al evento)',
         [('B', 'Desde (meses, incluido)'), ('C', 'Hasta (meses, sin llegar a)')])
    for i, m in enumerate(TRAMOS):
        f = F_TRA + 1 + i
        _f(ws, 'A%d' % f, 'Tramo %d' % (i + 1))
        _f(ws, 'B%d' % f, m, fmt=ENT, fill=VERDE, locked=False)
        _f(ws, 'C%d' % f, 'sin límite' if i == 0 else '=$B$%d' % (f - 1), fmt='General' if i == 0 else ENT)
    f = F_TRA + 1 + len(TRAMOS)
    _f(ws, 'A%d' % f, 'Sin huecos: cada tramo empieza donde termina el anterior y el último llega hasta el día del evento. Los '
                      'porcentajes retenidos son proporcionados y recíprocos: si cancela la finca, devuelve todo lo cobrado y '
                      'abona al cliente la misma penalización (art. 85.6 y 87.2 del TRLGDCU). Los porcentajes y los tramos son '
                      'ejemplos del kit: ajústalos a tu coste real de perder la fecha.', i=True, wrap=True)
    ws.merge_cells('A%d:F%d' % (f, f))
    ws.row_dimensions[f].height = 62
    # validaciones
    ref = refs('B')
    cel = lambda k: ref[k].replace('$', '')
    _dv(ws, 'list', [cel('canal')], f1='"%s"' % ','.join(CANALES), titulo='Cliente no válido', msg='Elige Particular o Empresa.')
    _dv(ws, 'list', [cel('tipo')], f1='=$A$%d:$A$%d' % (F_TIPOS + 1, F_TIPOS + len(TIPOS)), titulo='Tipología no válida',
        msg='Elige una tipología de la lista.')
    _dv(ws, 'date', [cel('fecha')], f1=str(serial(datetime.date(2026, 1, 1))), op='greaterThanOrEqual',
        titulo='Fecha no válida', msg='Escribe una fecha completa a partir de 2026 (dd/mm/aaaa).')
    _dv(ws, 'whole', [cel('guests')], f1='1', f2='100000', op='between', titulo='Número no válido',
        msg='Escribe un número entero de 1 o más.')
    _dv(ws, 'whole', [cel(k) for k in ('n_cer', 'n_hex', 'n_mon', 'n_alo')]
        + ['C%d' % (F_TEMP + i) for i in range(1, len(TEMPORADAS) + 1)]
        + ['B%d' % (F_TRA + 1 + i) for i in range(len(TRAMOS))] + ['%s%d' % (c, CT[k]) for k in ('d_2', 'd_3') for c in 'BC'],
        f1='0', f2='100000', op='between', titulo='Número no válido', msg='Escribe un número entero de 0 o más.')
    _dv(ws, 'decimal', [cel(k) for k in ('h_aux', 'wage')] + ['B%d' % X[k] for k in X]
        + ['%s%d' % (c, F_TIPOS + i) for c in 'BCD' for i in range(1, len(TIPOS) + 1)],
        f1='0', op='greaterThanOrEqual', titulo='Importe no válido', msg='Escribe un número mayor o igual que 0.')
    _dv(ws, 'decimal', [cel(k) for k in ('ss', 'cost_ext', 'canon', 'iva', 'iva_menu')]
        + ['%s%d' % (c, CT[k]) for k in CANAL_TAB if not k.startswith('d_') for c in 'BC'],
        f1='0', f2='1', op='between', titulo='Porcentaje no válido', msg='Escribe un porcentaje entre 0 y 1 (0,21 = 21 %).')
    _dv(ws, 'decimal', ['B%d' % (F_TEMP + i) for i in range(1, len(TEMPORADAS) + 1)]
        + ['B%d' % (F_DIA + i) for i in range(1, len(DIAS) + 1)], f1='-0.9', f2='2', op='between',
        titulo='Ajuste no válido', msg='Escribe un ajuste entre −90 % y +200 % (−0,20 = −20 %).')
    _dv(ws, 'list', ['B%d' % (F_MES + i) for i in range(1, len(MESES) + 1)], f1='"%s"' % ','.join(TEMPORADAS),
        titulo='Temporada no válida', msg='Elige Alta, Media o Baja.')
    _dv(ws, 'list', ['B%d' % (F_SEM + i) for i in range(1, len(SEMANA) + 1)], f1='"%s"' % ','.join(DIAS),
        titulo='Tipo de día no válido', msg='Elige Sábado, Viernes o domingo o Entre semana.')
    _dv(ws, 'whole', [cel('redondeo')], f1='1', f2='100000', op='between', titulo='Redondeo no válido',
        msg='Escribe un múltiplo de 1 euro o más (se divide entre él).')
    ws.freeze_panes = 'A5'
    _imprimir(ws)


def hoja_ejemplos(wb):
    ws = wb.create_sheet(H2)
    ws.column_dimensions['A'].width = 66
    for col in 'BCD':
        ws.column_dimensions[col].width = 28
    _f(ws, 'A1', '3 PRESUPUESTOS DE EJEMPLO — con las mismas fórmulas de la hoja «%s»' % H1, b=True, sz=14)
    _f(ws, 'A2', 'Estimaciones del kit para ilustrar el cálculo, no precios de mercado: cambia los datos en «%s» '
                 'para presupuestar tus eventos. Las tablas (tarifas, temporadas, días, extras, cobros y cancelación) están '
                 'al final de esa hoja y también mueven estos ejemplos.' % H1, wrap=True)
    ws.merge_cells('A2:D2')
    ws.row_dimensions[2].height = 32
    _f(ws, 'A4', 'DATOS DEL EVENTO', b=True, fill=GRIS)
    for j, nombre in enumerate(EJEMPLOS):
        _f(ws, '%s4' % 'BCD'[j], nombre, b=True, fill=GRIS, wrap=True)
    ws.row_dimensions[4].height = 48
    for k, rot, fmt, val, nota in ENTRADAS:
        _f(ws, 'A%d' % K[k], rot)
        for j, e in enumerate(EJEMPLOS.values()):
            _f(ws, '%s%d' % ('BCD'[j], K[k]), e[k], fmt=fmt if fmt != '@' else None, wrap=(fmt == '@'))
    _f(ws, 'A%d' % F_CALC, 'RESULTADO', b=True, fill=GRIS)
    for j in range(3):
        _f(ws, '%s%d' % ('BCD'[j], F_CALC), list(EJEMPLOS)[j], b=True, fill=GRIS, wrap=True)
    ws.row_dimensions[F_CALC].height = 48
    for k, rot, fmt, expr, nota in SALIDAS:
        neg = k in ('espacio', 'total', 'quoted', 'canon_est', 'margin_calc', 'presup')
        _f(ws, 'A%d' % S[k], rot, b=neg)
        for j in range(3):
            col = 'BCD'[j]
            _f(ws, '%s%d' % (col, S[k]), '=' + a_excel(expr, refs(col)), fmt=fmt, b=neg, wrap=(fmt == 'General'))
    f = S[SALIDAS[-1][0]] + 2
    _f(ws, 'A%d' % f, 'QUÉ ENSEÑA CADA EJEMPLO', b=True, fill=GRIS)
    for j, (nombre, nota) in enumerate(EJ_NOTAS.items()):
        _f(ws, 'A%d' % (f + 1 + j), nombre, b=True, wrap=True)
        _f(ws, 'B%d' % (f + 1 + j), nota, wrap=True)
        ws.merge_cells('B%d:D%d' % (f + 1 + j, f + 1 + j))
        ws.row_dimensions[f + 1 + j].height = 15 * max(3, math.ceil(len(nota) / 80.0))
    ws.freeze_panes = 'B5'
    _imprimir(ws, horizontal=True)


INSTRUCCIONES = [
    ('INSTRUCCIONES DE USO — Calculadora de presupuesto del espacio para bodas y eventos en finca', True),
    ('Convierte tu tarifa de alquiler en un presupuesto completo para un evento concreto: temporada y día de la fecha, '
     'mínimo de invitados, servicios extra, IVA, total, calendario de cobros, canon estimado del catering, margen del '
     'evento y lo que retienes si el cliente cancela. Sólo se escribe en las celdas VERDES; el margen se calcula.', False),
    ('1. Elige el cliente (Particular o Empresa), la tipología del evento, la fecha completa y los invitados previstos.', False),
    ('2. La fecha fija la temporada (tabla de meses) y el tipo de día (tabla de días de la semana). Cada una lleva un '
     'suplemento o un descuento visible en su tabla; los dos se multiplican con la tarifa de la tipología.', False),
    ('3. Cada temporada tiene un mínimo de invitados: se garantiza el mayor entre los previstos y el mínimo, y sobre esos '
     'invitados garantizados se estima el canon.', False),
    ('4. Escribe cuántos servicios extra llevas (ceremonia, horas extra, montaje premium, alojamiento). Sus precios están '
     'en la tabla «Servicios extra de la finca» y son de ejemplo: pon los tuyos.', False),
    ('5. Escribe las horas de personal de la finca, el coste bruto por hora, la Seguridad Social y el coste directo de los '
     'extras. El coste directo del evento (limpieza, aseos, gasoil, vigilancia y música) está en la tabla de tipologías.', False),
    ('IVA, CANON Y MENÚ', True),
    ('La finca factura el espacio y los servicios con precio propio al 21 %. El menú y la bebida los factura el catering '
     'exclusivo: comida y bebida para consumir en el acto van al 10 % (Ley 37/1992, art. 91.Uno.2.2.º), y no son ingreso de '
     'la finca. La calculadora los muestra sólo como orientación para el cliente. Si vendes un paquete que mezcla ambos, '
     'consúltalo con tu asesoría.', False),
    ('El canon es lo que el catering exclusivo paga a la finca por tener la exclusiva. Aquí es un supuesto editable (12 % de '
     'lo que factura el catering): no es un dato de mercado, negócialo. Lo paga el catering, no el cliente; la finca lo '
     'factura con IVA.', False),
    ('El margen del evento es una contribución: no descuenta la renta, el personal fijo, los suministros, los seguros ni '
     'la estructura. Eso está en el plan financiero.', False),
    ('COBRO Y CANCELACIÓN', True),
    ('Cobro en tres pagos: reserva al firmar, segundo pago y liquidación antes del evento, con sus fechas límite. Emite '
     'factura de cada pago, también de la reserva. Particular: 25 %, 35 % a seis meses y 40 % a 30 días. Empresa: 30 %, 30 % '
     'a dos meses y 40 % a 15 días. Son ejemplos del kit.', False),
    ('La cancelación se escribe por tramos en meses, sin huecos: el primero empieza a los 12 meses o más y el último llega '
     'hasta el día del evento. Los porcentajes retenidos suben con la cercanía y son proporcionados al coste de perder la fecha '
     '(art. 85.6 del TRLGDCU). Tienen que ser recíprocos: si cancela la finca, devuelve lo cobrado y paga al cliente la misma '
     'penalización; retener lo pagado sin prever esa indemnización es abusivo (art. 87.2).', False),
    ('Un servicio de celebración con fecha fijada puede quedar fuera del derecho de desistimiento de la contratación a '
     'distancia (art. 103.l del TRLGDCU): por eso la política de cancelación tiene que estar firmada en el contrato.', False),
    ('ESTIMACIONES DEL KIT', True),
    ('Los valores por defecto y los tres ejemplos son estimaciones para ilustrar el cálculo, no precios de mercado. '
     'Cámbialos por tus tarifas, tus extras, tus costes y tus tramos antes de presupuestar.', False),
]


def hoja_instrucciones(wb):
    ws = wb.create_sheet(H3)
    ws.column_dimensions['A'].width = 110
    f = 1
    for t, b in INSTRUCCIONES:
        _f(ws, 'A%d' % f, t, b=b, sz=13 if f == 1 else 11, wrap=True)
        if not b:
            ws.row_dimensions[f].height = max(16, 15 * math.ceil(len(t) / 120.0))
        f += 2 if b and f > 1 else 1
    f += 1
    for t in ('Para editar una celda que no esté en verde: Revisar → Desproteger hoja (no tiene contraseña).',
              'Más plantillas y kits del catálogo en aichef.pro/productos-digitales', motor.BIO_LINE,
              motor.VERSION_FMT.format(pid=PID)):
        _f(ws, 'A%d' % f, t, wrap=True)
        f += 1
    _imprimir(ws)


def generar(destino):
    wb = openpyxl.Workbook()
    hoja_calculadora(wb)
    hoja_ejemplos(wb)
    hoja_instrucciones(wb)
    pr = wb.properties
    pr.title = '%s · v%s' % (TITULO, motor.VERSION)
    pr.subject = 'Plan de Negocio: Finca de Eventos y Bodas'
    pr.description = 'aichef.pro/%s' % PID
    pr.keywords = 'plan negocio finca eventos bodas, AI Chef Pro'
    pr.creator = pr.lastModifiedBy = motor.CREATOR
    os.makedirs(os.path.dirname(destino), exist_ok=True)
    wb.save(destino)
    sys.path.insert(0, PD)
    import inject_cache as IC                                                      # noqa: E402
    nf, inj, fail = IC.inject(destino)
    print('%s: formulas=%d cache_inyectado=%d fallos_pycel=%d' % (os.path.basename(destino), nf, inj, fail))
    return fail


PROHIBIDOS = [r'\bfreelance', r'validad[oa]s?', r'Espa[ñn]a,? 2026', r'food ?truck', r'carnet', r'carné',
              r'tarifa plana', r'cocina central', r'parrill', r'RD 393', r'carné de manipulador']


def _num(x):
    """Valor leído de la caché → número (las fechas vuelven como datetime)."""
    if isinstance(x, datetime.datetime):
        return serial(x.date())
    if isinstance(x, datetime.date):
        return serial(x)
    return x


def _distinto(x, v):
    x = _num(x)
    if isinstance(v, str):
        return x != v
    return not isinstance(x, (int, float)) or abs(float(x) - v) > 0.005


def autotest(destino):
    from pycel import ExcelCompiler
    err, n = [], 0
    xl = ExcelCompiler(filename=destino)
    wv = openpyxl.load_workbook(destino, data_only=True)
    wf = openpyxl.load_workbook(destino)
    casos = [(H2, 'BCD'[j], e, nombre) for j, (nombre, e) in enumerate(EJEMPLOS.items())]
    casos.append((H1, 'B', OrderedDict((k, v) for k, _r2, _f2, v, _n in ENTRADAS), 'hoja principal'))
    for hoja, col, e, nombre in casos:
        esperado = calcular(e)
        for k, v in esperado.items():
            c = '%s%d' % (col, S[k])
            n += 1
            for lab, x in (('pycel', xl.evaluate("'%s'!%s" % (hoja, c))), ('caché', wv[hoja][c].value)):
                if _distinto(x, v):
                    err.append('%s · %s %s (%s): %s %r ≠ esperado %r' % (nombre, hoja, c, k, lab, x, v))
        if not 10 <= esperado['eur_inv'] <= 80:
            err.append('%s: espacio y servicios %.2f €/inv fuera de 10-80' % (nombre, esperado['eur_inv']))
    # los ejemplos enseñan lo que dicen
    pe = [calcular(e) for e in EJEMPLOS.values()]
    if (pe[0]['temporada'], pe[0]['dia']) != ('Alta', 'Sábado') or (pe[1]['temporada'], pe[1]['dia']) != ('Media', 'Viernes o domingo') \
            or (pe[2]['temporada'], pe[2]['dia']) != ('Media', 'Entre semana') or pe[2]['garant'] <= EJEMPLOS['Jornada de empresa entre semana']['guests']:
        err.append('los ejemplos no son sábado alta / viernes media / entre semana con mínimo aplicado')
    # casos propios: cambian celdas de la hoja principal y se recalculan con pycel
    base = OrderedDict((k, v) for k, _r2, _f2, v, _n in ENTRADAS)
    propios = [('propio 1: Empresa, lunes de enero, 20 invitados, celebración pequeña',
                dict(base, canal='Empresa', tipo='Celebración pequeña', fecha=serial(datetime.date(2027, 1, 18)), guests=20,
                     n_cer=0, n_hex=0, n_mon=0, h_aux=4, redondeo=10)),
               ('propio 2: Particular, domingo de septiembre, 150 invitados, celebración privada',
                dict(base, canal='Particular', tipo='Celebración privada', fecha=serial(datetime.date(2027, 9, 12)), guests=150,
                     n_cer=2, n_hex=3, n_mon=2, n_alo=6, canon=0.10, h_aux=20)),
               ('propio 3: Empresa, sábado de agosto, 85 invitados, evento de empresa',
                dict(base, canal='Empresa', tipo='Empresa', fecha=serial(datetime.date(2027, 8, 7)), guests=85,
                     n_cer=0, n_hex=4, n_mon=1, n_alo=0, iva=0.21, iva_menu=0.10))]
    for nombre, e in propios:
        esperado = calcular(e)
        xl2 = ExcelCompiler(filename=destino)
        for k in S:
            xl2.evaluate("'%s'!B%d" % (H1, S[k]))
        for k, v in e.items():
            xl2.set_value("'%s'!B%d" % (H1, K[k]), v)
        for k, v in esperado.items():
            n += 1
            x = xl2.evaluate("'%s'!B%d" % (H1, S[k]))
            if _distinto(x, v):
                err.append('%s: %s pycel %r ≠ esperado %r' % (nombre, k, x, v))
    # el selector de cliente cambia el cobro y la cancelación
    e = dict(EJEMPLOS['Boda de sábado en temporada alta'])
    cs = [calcular(dict(e, canal=c)) for c in CANALES]
    if cs[0]['cb1'] == cs[1]['cb1'] or cs[0]['can2'] == cs[1]['can2'] or cs[0]['quoted'] == cs[1]['quoted']:
        err.append('el selector de cliente no cambia el cobro, la cancelación o el importe presupuestado')
    # cobros suman el 100 %, cancelación creciente, proporcionada y sin huecos
    for ic, canal in enumerate(CANALES):
        s = sum(CANAL_TAB[k][1][ic] for k in ('pc_1', 'pc_2', 'pc_3'))
        if abs(s - 1.0) > 1e-9:
            err.append('cobro %s suma %.2f ≠ 100 %%' % (canal, s))
        v = [CANAL_TAB['cn%d' % i][1][ic] for i in range(1, 6)]
        if v != sorted(v) or v[-1] != 1.0:
            err.append('cancelación %s: tramos no crecientes o el último no es el 100 %% %r' % (canal, v))
        if v[0] > CANAL_TAB['pc_1'][1][ic]:
            err.append('cancelación %s: el tramo 1 retiene más que la reserva' % canal)
        if CANAL_TAB['d_2'][1][ic] <= CANAL_TAB['d_3'][1][ic]:
            err.append('cobro %s: el segundo pago no es anterior a la liquidación' % canal)
    if TRAMOS != sorted(TRAMOS, reverse=True) or len(set(TRAMOS)) != len(TRAMOS) or TRAMOS[-1] != 0 or len(TRAMOS) != 5:
        err.append('tramos de cancelación con huecos o desordenados: %r' % TRAMOS)
    ws = wf[H1]
    for i in range(len(TRAMOS)):          # la columna «hasta» enlaza con el «desde» del tramo anterior (sin huecos)
        v = ws['C%d' % (F_TRA + 1 + i)].value
        if i and v != '=$B$%d' % (F_TRA + i):
            err.append('tramo %d: «hasta» no enlaza con el «desde» anterior (%r)' % (i + 1, v))
    # cada mes y cada día de la semana cae en una temporada / tipo de día de la tabla
    if any(t not in TEMPORADAS for _m, t in MESES) or any(d not in DIAS for _s, d in SEMANA) or len(MESES) != 12 or len(SEMANA) != 7:
        err.append('tabla de meses o de días de la semana incompleta')
    if motor.detectar(wf, FICHERO)['molde'] != 'calc-beta':
        err.append('motor.detectar no la reconoce como calc-beta')
    gf = motor.gate_formatos(wf, FICHERO)
    if gf:
        err.append('gate_formatos: %d celdas, p. ej. %r' % (len(gf), gf[:3]))
    # validaciones de datos
    dvs = dict((str(rng), dv) for dv in ws.data_validations.dataValidation for rng in dv.sqref.ranges)
    c = 'B%d' % K['canal']
    if dvs.get(c) is None or dvs[c].type != 'list' or 'Particular,Empresa' not in (dvs[c].formula1 or ''):
        err.append('DV de lista incorrecta en %s' % c)
    c = 'B%d' % K['tipo']
    if dvs.get(c) is None or dvs[c].type != 'list' or not (dvs[c].formula1 or '').startswith('=$A$'):
        err.append('DV de lista del rango incorrecta en %s' % c)
    for k, tp in (('fecha', 'date'), ('guests', 'whole'), ('n_cer', 'whole')):
        dv = dvs.get('B%d' % K[k])
        if dv is None or dv.type != tp:
            err.append('DV %s ausente en %s' % (tp, k))
    textos = []
    for w in wf.worksheets:
        if w.page_setup.paperSize != 9 or not w.protection.sheet:
            err.append('%s: A4 y protección' % w.title)
        for row in w.iter_rows():
            for c in row:
                v = c.value
                if isinstance(v, str) and c.data_type != 'f':
                    textos.append((w.title, c.coordinate, v))
                if c.data_type == 'f' and wv[w.title][c.coordinate].value is None:
                    err.append('%s!%s: fórmula sin caché' % (w.title, c.coordinate))
                if c.fill.fgColor.rgb in ('00E8F5E9', 'FFE8F5E9') and c.protection.locked:
                    err.append('%s!%s: verde bloqueada' % (w.title, c.coordinate))
    # ninguna tasa tecleada dentro de una fórmula: sólo constantes estructurales (0, 1)
    for w in wf.worksheets:
        for row in w.iter_rows():
            for c in row:
                if c.data_type == 'f':
                    sin_refs = re.sub(r"'[^']+'!|\$?[A-Z]{1,2}\$?\d+", '', str(c.value))
                    for num in re.findall(r'(?<![A-Za-z])\d+\.?\d*', sin_refs):
                        if num not in ('0', '1'):
                            err.append('%s!%s: número tecleado %s en la fórmula %s' % (w.title, c.coordinate, num, c.value))
    todo = '\n'.join(t for _h, _c, t in textos)
    for rx in PROHIBIDOS:
        if re.search(rx, todo, re.I):
            err.append('prohibido /%s/' % rx)
    if re.search(r'service charge|sales tax|\$', todo, re.I):
        err.append('restos del EN (service charge / sales tax / $)')
    if re.search(r'[^\x00-ɏ‐-‧€−→·]', todo):
        err.append('caracteres no latinos o raros: %r' % sorted(set(re.findall(r'[^\x00-ɏ‐-‧€−→·]', todo))))
    if motor.VERSION_FMT.format(pid=PID) not in todo or motor.BIO_LINE not in todo:
        err.append('falta la línea de versión o la bio')
    if '1.0' not in motor.VERSION_FMT or 'noviembre 2026' not in motor.VERSION_FMT:
        err.append('versión no es 1.0 · noviembre 2026: %s' % motor.VERSION_FMT)
    if wf.sheetnames != [H1, H2, H3]:
        err.append('pestañas %r' % wf.sheetnames)
    p = wf.properties
    if p.creator != 'AI Chef Pro' or p.description != 'aichef.pro/%s' % PID or \
            p.keywords != 'plan negocio finca eventos bodas, AI Chef Pro':
        err.append('docProps: creator=%r description=%r keywords=%r' % (p.creator, p.description, p.keywords))
    for x in err:
        print('ERROR', x)
    print('autotest calculadora finca ES: %d comprobaciones (pycel + caché en ejemplos y hoja principal, pycel en '
          '3 casos propios) · %s' % (n * 2, 'VERDE' if not err else 'ROJO (%d errores)' % len(err)))
    for nombre, e in EJEMPLOS.items():
        r = calcular(e)
        print('  %-36s %s/%s · espacio %8.2f · base %8.2f · total IVA %8.2f · reserva %8.2f · canon %7.2f · margen %5.1f %%'
              % (nombre, r['temporada'], r['dia'], r['espacio'], r['base'], r['total'], r['cb1'], r['canon_est'],
                 r['margin_calc'] * 100))
    return 1 if err else 0


if __name__ == '__main__':
    ap = argparse.ArgumentParser()
    ap.add_argument('--salida', required=True, help='carpeta de salida')
    ap.add_argument('--autotest', action='store_true')
    a = ap.parse_args()
    destino = os.path.join(a.salida, FICHERO)
    if not a.autotest and generar(destino):
        sys.exit('ERROR: fallos de pycel en inject_cache')
    sys.exit(autotest(destino))
