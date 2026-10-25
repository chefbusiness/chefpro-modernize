#!/usr/bin/env python3
"""barrido_es.py — Finca de Eventos y Bodas ES 1.0 (F1-camino §3; copia de `plan-heladeria/barrido_es.py`): barre los
xlsx/docx que se le pasen buscando
  · restos de los TRES moldes: panadería (semilla: pan, bollería, harina, masa madre, hornada, levadura, madrugada,
    transacción…), heladería (helado, mantecadora, cubeta, vitrina, obrador, cucurucho, tarrina) y parrillero/catering
    (cocina de apoyo, furgoneta, remolque, parrilla, brasa, carbón, asado para llevar, cocina central, drop-off), más el
    Food Truck (food truck, camión, ventanilla);
  · frases del motor sin traducir por el VOCABULARIO (clientes/día, días de apertura, el local, en sala, ticket, delivery,
    cubierto, comensal);
  · los prohibidos de la SPEC D11 (RD 393/2007 derogado, «carné de manipulador», «España 2026», el 10 % para todo el
    paquete, el canon del 12 % como dato sin «negócialo»…) y los de la familia (freelance, validado, tarifa plana, CNAE,
    licencia de apertura, versión anterior…);
  · concordancias de «la finca» (femenino) que el motor escribía para «el local».
SIN «generador» ni «aparcamiento» (legítimos aquí: grupo electrógeno y aparcamiento de invitados). Lista celda/párrafo
y texto; en las fórmulas sólo mira los literales entre comillas. `PROHIBIDOS` se puede importar desde `tandas_docx.py`.
Corre en el VPS. Uso: barrido_es.py <fichero>… · exit 1 si hay coincidencias. Sesión Claude Code, 25-oct-2026."""
import re
import sys

PROHIBIDOS = [
    # familia y SPEC D11
    r'\bfreelance', r'validad[oa]s?', r'Espa[ñn]a,? 2026', r'carn[eé]t? de manipulador(?! ya no existe)',
    r'tarifa plana', r'\bCNAE\b', r'(?<!inscrito en el )\bRGSEAA\b(?!.{0,80}(s[oó]lo|otros (establecimientos|negocios)))',
    r'licencia de apertura', r'versi[oó]n (1\.1|2\.2|anterior)', r'\bv1\.1\b', r'solopreneur',
    r'393/2007', r'\bRD 393\b', r'(todo el|el) (paquete|pack)[^.;]{0,60}\b10 ?%', r'\b10 ?%[^.;]{0,40}(todo el|el) (paquete|pack)',
    r'canon (del|de un) 12 ?%(?![^.]{0,120}(neg[oó]ci|supuesto))',
    # Food Truck
    r'food ?truck', r'\bcami[oó]n\b', r'ventanilla',
    # panadería (semilla)
    r'panader', r'\bpan\b', r'boller', r'harina', r'masa madre', r'hornada', r'horno de pisos', r'amasad',
    r'levadura', r'madrugada', r'pan com[uú]n', r'superreducido', r'transacci',
    # heladería (semilla clonada)
    r'helad', r'mantec', r'cubeta', r'vitrina', r'obrador', r'cucurucho', r'tarrina',
    # parrillero y catering (a.py clonado)
    r'cocina de apoyo', r'furgoneta', r'remolque', r'parrill', r'\bbrasa', r'carb[oó]n', r'asado para llevar',
    r'cocina central', r'drop-?off',
    # motor sin traducir
    r'\bclientes?/d[ií]a', r'\bd[ií]as de apertura', r'\bdel local\b', r'\bel local\b', r'\bde local\b',
    r'\ben sala\b', r'\bticket\b', r'\bdelivery\b', r'\bcubiert', r'\bcomensal',
    # concordancias de «la finca»
    r'\b(el|del|al|un|este|ese|nuestro|vuestro) finca\b',
    r'\bfinca (nuevo|abierto|cerrado|arrendado|alquilado|propio|vac[ií]o|lleno|situado|ubicado)\b',
]
RX = re.compile('|'.join(PROHIBIDOS), re.I)


def textos(p):
    import docx
    import openpyxl
    if p.endswith('.xlsx'):
        wb = openpyxl.load_workbook(p)
        it = [(ws.title + '!' + c.coordinate, c.value) for ws in wb.worksheets for row in ws.iter_rows() for c in row
              if isinstance(c.value, str)]
        it += [(ws.title + '!pestaña', ws.title) for ws in wb.worksheets]
        pr = wb.properties
        it += [('docProps.' + k, getattr(pr, k)) for k in ('title', 'subject', 'description', 'keywords', 'creator')
               if isinstance(getattr(pr, k), str)]
    else:
        d = docx.Document(p)
        it = [('p%d' % i, x.text) for i, x in enumerate(d.paragraphs)]
        it += [('tabla%d' % i, c.text) for i, t in enumerate(d.tables) for row in t.rows for c in row.cells]
        cp = d.core_properties
        it += [('docProps.' + k, getattr(cp, k)) for k in ('title', 'subject', 'comments', 'keywords', 'author')
               if isinstance(getattr(cp, k), str)]
    return it


def main(rutas):
    n = 0
    for p in [x for x in rutas if not x.startswith('--')]:
        for loc, t in textos(p):
            txt = ' '.join(re.findall(r'"([^"]*)"', t)) if t.startswith('=') else t
            m = RX.search(txt)
            if m:
                n += 1
                print('%s · %s · «%s» · …%s…' % (p.split('/')[-1][:28], loc, m.group(0),
                                               txt[max(0, m.start() - 50):m.end() + 50]))
    print('barrido: %d coincidencias' % n)
    return n


if __name__ == '__main__':
    sys.exit(1 if main(sys.argv[1:]) else 0)
