#!/usr/bin/env python3
"""
cifras_fin_es.py — Finca de Eventos y Bodas ES 1.0 (copia de `plan-heladeria/cifras_hel_es.py` con el MAPA de eventos de
`parrillero-2_2/cifras_par_es.py`; rótulos de línea A medidos sobre el libro generado el 25-oct-2026, ya pasados por el
VOCABULARIO de `contenido_plan_negocio_finca_eventos_bodas/a.py`): lee la CACHÉ del plan financiero (después de
`inject_cache`, que hace `main.py`) por RÓTULO, escribe `cifras_caso_es.json` (los tokens {ASI} del Word) y pasa los
gates del caso de la SPEC D9:
  · DSCR mínimo ≥ 1,30 · saldo mínimo de tesorería > 0 · cobertura de horas ≥ 98 %;
  · holgura sobre el equilibrio de caja del año 2 ≥ 12 % y del año 3 ≥ 15 % (la del año 1, informativa);
  · resultado neto del año 3 entre el 6 y el 14 % de las ventas (el del año 1, informativo);
  · exceso de financiación (origen − usos) ≥ 0 y < 5 % (el mismo control que la demo RD-34 del motor);
  · D4: las dos líneas de venta al 21 % (columna G del P&L) y el IVA repercutido de la Tesorería = 21 % de las ventas
    de cada mes; el resultado de cada liquidación trimestral (303) = repercutido − soportado − lo arrastrado.
Corre en el VPS (openpyxl).

    /root/venv-guias/bin/python cifras_fin_es.py --xlsx <plan-financiero-finca-eventos-bodas.xlsx>

Sesión Claude Code, 25-oct-2026.
"""
import argparse
import json
import math
import os
import re
import sys

import openpyxl

AQUI = os.path.dirname(os.path.abspath(__file__))
SALIDA = os.path.join(AQUI, 'cifras_caso_es.json')

#: Mes 1 = marzo (SPEC D8): nombre del mes del calendario de cada «Mes n» de la Tesorería.
MESES = ('marzo', 'abril', 'mayo', 'junio', 'julio', 'agosto', 'septiembre', 'octubre', 'noviembre', 'diciembre',
         'enero', 'febrero')

#: token → (hoja, regex del rótulo de la columna A, columna, formato). Formatos: eur (miles con punto, sin
#: decimales) · eur2 · pct (1 decimal) · pct0 · dec1 · dec2 · int · int_up (redondeo hacia arriba).
SUP, PYG, INV, PE, TES, FIN, PER, ESC = ('0. Supuestos', 'PyG 3 Años', 'Inversión Inicial', 'Punto Equilibrio',
                                         'Tesorería 12 meses', 'Financiación', 'Personal', 'Escenarios')
MAPA = {
    # Supuestos (driver sin permutar: invitados por evento × precio por invitado × eventos al año)
    'INVITADOS_EVENTO': (SUP, r'^invitados por evento \(media', 2, 'int'),
    'PRECIO_INVITADO': (SUP, r'^precio medio por invitado sin iva', 2, 'eur2'),
    'PRECIO_INVITADO_IVA': (SUP, r'^pvp equivalente con iva', 2, 'eur2'),
    'EVENTOS_ANO': (SUP, r'^eventos al a[ñn]o$', 2, 'int'),
    'CREC_A2': (SUP, r'^crecimiento de volumen del a[ñn]o 2', 2, 'pct0'),
    'CREC_A3': (SUP, r'^crecimiento de volumen del a[ñn]o 3', 2, 'pct0'),
    'PCT_ESPACIO': (SUP, r'^ventas de espacio y servicios sobre el total', 2, 'pct0'),
    'PCT_CANON': (SUP, r'^ventas del canon del catering sobre el total', 2, 'pct0'),
    'COSTE_DIRECTO_PCT': (SUP, r'^coste directo del evento sobre las ventas', 2, 'pct'),
    'CONSUMIBLES_PCT': (SUP, r'^consumibles sobre ventas', 2, 'pct0'),
    'INTERMEDIARIOS_PCT': (SUP, r'^ventas que llegan por wedding planners', 2, 'pct0'),
    'COMISION_INTERMEDIARIOS': (SUP, r'^comisi[oó]n de los wedding planners', 2, 'pct'),
    'VARIOS_PCT': (SUP, r'^varios e imprevistos sobre ventas', 2, 'pct0'),
    'RENTA_MES': (SUP, r'^renta mensual de la finca', 2, 'eur'),
    'FIANZA_MESES': (SUP, r'^fianza del arrendamiento \(meses', 2, 'int'),
    'SUMINISTROS_MES': (SUP, r'^suministros mensuales', 2, 'eur'),
    'SEGUROS_ANO': (SUP, r'^seguros', 2, 'eur'),
    'RECURSOS_PROPIOS': (SUP, r'^recursos propios', 2, 'eur'),
    'TIPO_PRESTAMO': (SUP, r'^tipo de inter[eé]s nominal', 2, 'pct'),
    'PLAZO_PRESTAMO': (SUP, r'^plazo del pr[eé]stamo', 2, 'int'),
    'CARENCIA': (SUP, r'^carencia de principal', 2, 'int'),
    'MESES_FONDO': (SUP, r'^fondo de maniobra \(meses', 2, 'int'),
    'VIDA_OBRA': (SUP, r'^vida [uú]til de obra', 2, 'int'),
    'VIDA_MAQUINARIA': (SUP, r'^vida [uú]til de maquinaria', 2, 'int'),
    'DIAS_COBRO': (SUP, r'^d[ií]as medios de cobro', 2, 'int'),
    'AFORO': (SUP, r'^aforo de la finca', 2, 'int'),
    'RAMPA_MES1': (SUP, r'^actividad del mes 1', 2, 'pct0'),
    'MESES_RAMPA': (SUP, r'^meses hasta alcanzar', 2, 'int'),
    'IS_GENERAL': (SUP, r'^impuesto de sociedades, tipo general', 2, 'pct0'),
    'IS_NUEVA': (SUP, r'^impuesto de sociedades, nueva creaci', 2, 'pct0'),
    'HOLGURA_MINIMA': (SUP, r'^holgura m[ií]nima sobre el punto de equilibrio', 2, 'pct0'),
    # P&L
    'VENTAS_A1': (PYG, r'^ingresos totales', 2, 'eur'),
    'VENTAS_A2': (PYG, r'^ingresos totales', 3, 'eur'),
    'VENTAS_A3': (PYG, r'^ingresos totales', 4, 'eur'),
    'INVITADOS_EVENTO_A2': (PYG, r'^invitados por evento \(media\)', 3, 'dec1'),
    'INVITADOS_EVENTO_A3': (PYG, r'^invitados por evento \(media\)', 4, 'dec1'),
    'VENTAS_ESPACIO_A1': (PYG, r'^ventas de alquiler del espacio', 2, 'eur'),
    'CANON_A1': (PYG, r'^canon del catering exclusivo', 2, 'eur'),
    'CANON_A3': (PYG, r'^canon del catering exclusivo', 4, 'eur'),
    'IVA_LINEA_ESPACIO': (PYG, r'^ventas de alquiler del espacio', 7, 'pct0'),
    'IVA_LINEA_CANON': (PYG, r'^canon del catering exclusivo', 7, 'pct0'),
    'COSTE_DIRECTO_A1': (PYG, r'^coste directo del evento$', 2, 'eur'),
    'COMISIONES_INTERMEDIARIOS_A1': (PYG, r'^comisiones de wedding planners', 2, 'eur'),
    'COSTES_VARIABLES_A1': (PYG, r'^total costes variables', 2, 'eur'),
    'MARGEN_BRUTO_A1': (PYG, r'^margen bruto$', 2, 'eur'),
    'MARGEN_BRUTO_PCT_A1': (PYG, r'^margen bruto / ventas', 2, 'pct'),
    'RENTA_ANO': (PYG, r'^renta de la finca$', 2, 'eur'),
    'RENTA_PCT_A1': (PYG, r'^renta de la finca / ventas', 2, 'pct'),
    'PERSONAL_A1': (PYG, r'^personal \(n[oó]minas', 2, 'eur'),
    'PERSONAL_A2': (PYG, r'^personal \(n[oó]minas', 3, 'eur'),
    'PERSONAL_A3': (PYG, r'^personal \(n[oó]minas', 4, 'eur'),
    'PERSONAL_PCT_A1': (PYG, r'^coste de personal / ventas', 2, 'pct'),
    'COSTE_DIRECTO_PCT_A1': (PYG, r'^coste directo del evento / ventas', 2, 'pct'),
    'SUMINISTROS_ANO': (PYG, r'^suministros \(', 2, 'eur'),
    'MARKETING_ANO': (PYG, r'^marketing: portales de bodas', 2, 'eur'),
    'FERIAS_ANO': (PYG, r'^ferias nupciales', 2, 'eur'),
    'JARDINERIA_ANO': (PYG, r'^jardiner[ií]a y mantenimiento externo', 2, 'eur'),
    'AMORTIZACION_A1': (PYG, r'^amortizaci[oó]n del inmovilizado', 2, 'eur'),
    'INTERESES_A1': (PYG, r'^gastos financieros', 2, 'eur'),
    'FIJOS_A1': (PYG, r'^total costes fijos', 2, 'eur'),
    'RAI_A1': (PYG, r'^resultado antes de impuestos', 2, 'eur'),
    'RAI_A2': (PYG, r'^resultado antes de impuestos', 3, 'eur'),
    'RAI_A3': (PYG, r'^resultado antes de impuestos', 4, 'eur'),
    'IS_A1': (PYG, r'^impuesto de sociedades$', 2, 'eur'),
    'NETO_A1': (PYG, r'^resultado neto$', 2, 'eur'),
    'NETO_A2': (PYG, r'^resultado neto$', 3, 'eur'),
    'NETO_A3': (PYG, r'^resultado neto$', 4, 'eur'),
    'NETO_PCT_A1': (PYG, r'^resultado neto / ventas', 2, 'pct'),
    'NETO_PCT_A2': (PYG, r'^resultado neto / ventas', 3, 'pct'),
    'NETO_PCT_A3': (PYG, r'^resultado neto / ventas', 4, 'pct'),
    # Inversión
    'CAPEX': (INV, r'^capex', 2, 'eur'),
    'FONDO_MANIOBRA': (INV, r'^colch[oó]n operativo', 2, 'eur'),
    'INVERSION_TOTAL': (INV, r'^inversi[oó]n total', 2, 'eur'),
    'IVA_INVERSION': (INV, r'^iva soportado sobre la inversi', 2, 'eur'),
    'NECESIDAD_CAJA': (INV, r'^necesidad total de caja', 2, 'eur'),
    'FIANZA': (INV, r'^fianza del arrendamiento', 2, 'eur'),
    'RENTA_PREVIA': (INV, r'^renta de la finca durante la obra', 2, 'eur'),
    'IMPREVISTOS': (INV, r'^imprevistos de obra', 2, 'eur'),
    'PROYECTO': (INV, r'^proyecto t[eé]cnico', 2, 'eur'),
    'LICENCIAS': (INV, r'^licencias, tasas', 2, 'eur'),          # revisión final M8: el rótulo empieza por «Licencias»
    'SALON': (INV, r'^adecuaci[oó]n del sal[oó]n', 2, 'eur'),
    'ACUSTICA': (INV, r'^ac[uú]stica', 2, 'eur'),
    'CARPA': (INV, r'^carpa', 2, 'eur'),
    'ASEOS': (INV, r'^aseos', 2, 'eur'),
    'ELECTRICA': (INV, r'^instalaci[oó]n el[eé]ctrica', 2, 'eur'),
    'GRUPO_ELECTROGENO': (INV, r'^grupo electr[oó]geno', 2, 'eur'),
    'CLIMATIZACION': (INV, r'^climatizaci[oó]n', 2, 'eur'),
    'OFFICE': (INV, r'^office del catering', 2, 'eur'),
    'MOBILIARIO': (INV, r'^mobiliario de ceremonia', 2, 'eur'),
    'SONIDO': (INV, r'^sonido', 2, 'eur'),
    'APARCAMIENTO': (INV, r'^aparcamiento y caminos', 2, 'eur'),
    'JARDIN': (INV, r'^jardiner[ií]a, paisajismo', 2, 'eur'),
    'PCI': (INV, r'^protecci[oó]n contra incendios', 2, 'eur'),
    'LANZAMIENTO': (INV, r'^lanzamiento', 2, 'eur'),
    'BASE_OBRA': (INV, r'^obra, instalaciones y acondicionamiento$', 2, 'eur'),
    'BASE_MAQUINARIA': (INV, r'^maquinaria, mobiliario y equipos', 2, 'eur'),
    # Financiación
    'PRESTAMO': (FIN, r'^pr[eé]stamo bancario$', 2, 'eur'),
    'ORIGEN_FONDOS': (FIN, r'^total origen de fondos', 2, 'eur'),
    'DIFERENCIA_FONDOS': (FIN, r'^diferencia \(origen - usos\)', 2, 'eur'),
    'DIFERENCIA_FONDOS_PCT': (FIN, r'^diferencia sobre los usos', 2, 'pct'),
    'CUOTA_ANUAL': (FIN, r'^cuota anual durante la amortizaci', 2, 'eur'),
    'DSCR_A2': (FIN, r'^dscr del a[ñn]o 2', 2, 'dec2'),
    'DSCR_A3': (FIN, r'^dscr del a[ñn]o 3', 2, 'dec2'),
    'DSCR_MIN': (FIN, r'^dscr m[ií]nimo de todo el cuadro', 2, 'dec2'),
    # Tesorería
    'SALDO_MINIMO': (TES, r'^saldo m[ií]nimo del a[ñn]o', 2, 'eur'),
    'MES_SALDO_MINIMO': (TES, r'^mes en el que la caja toca fondo', 2, 'int'),
    'SALDO_MES12': (TES, r'^saldo acumulado de caja', 13, 'eur'),
    'IVA_REPERCUTIDO_A1': (TES, r'^iva repercutido del mes', 14, 'eur'),
    'PAGO_303_A1': (TES, r'^pago del iva \(modelo 303\)', 14, 'eur'),
    'PAYBACK': (TES, r'^payback del proyecto', 2, 'dec1'),
    'PAYBACK_CAPEX': (TES, r'^payback sobre el capex', 2, 'dec1'),
    # Punto de equilibrio
    'COSTE_VAR_INVITADO': (PE, r'^coste variable por invitado', 2, 'eur2'),
    'MARGEN_INVITADO': (PE, r'^margen de contribuci[oó]n por invitado', 2, 'eur2'),
    'BE_INVITADOS_EVENTO': (PE, r'^invitados necesarios por evento$', 2, 'int_up'),
    'BE_INVITADOS_EVENTO_CAJA': (PE, r'^invitados necesarios por evento \(caja\)', 2, 'int_up'),
    'BE_INVITADOS_EVENTO_CAJA_A2': (PE, r'^invitados necesarios por evento \(caja\)', 3, 'int_up'),
    'BE_INVITADOS_EVENTO_CAJA_A3': (PE, r'^invitados necesarios por evento \(caja\)', 4, 'int_up'),
    'BE_INGRESOS': (PE, r'^ingresos necesarios al a[ñn]o$', 2, 'eur'),
    'BE_INGRESOS_CAJA': (PE, r'^ingresos necesarios al a[ñn]o \(caja\)', 2, 'eur'),
    'BE_INGRESOS_CAJA_A2': (PE, r'^ingresos necesarios al a[ñn]o \(caja\)', 3, 'eur'),
    'BE_INGRESOS_CAJA_A3': (PE, r'^ingresos necesarios al a[ñn]o \(caja\)', 4, 'eur'),
    'HOLGURA_CAJA': (PE, r'^holgura sobre el equilibrio de caja', 2, 'pct'),
    'HOLGURA_CAJA_A2': (PE, r'^holgura sobre el equilibrio de caja', 3, 'pct'),
    'HOLGURA_CAJA_A3': (PE, r'^holgura sobre el equilibrio de caja', 4, 'pct'),
    # Personal
    'COBERTURA_HORAS': (PER, r'^cobertura \(contratadas', 2, 'pct0'),
    'PLANTILLA_PERSONAS': (PER, r'^total plantilla', 2, 'int'),
    'SUELDO_TITULAR': (PER, r'^titular', 4, 'eur'),
    'HORAS_EVENTO': (PER, r'^horas de trabajo por evento', 2, 'dec1'),
    'PERSONAS_EVENTO': (PER, r'^personas necesarias a la vez en cada evento', 2, 'dec1'),
    'HORAS_ESTRUCTURA': (PER, r'^horas de estructura por evento', 2, 'int'),
    'HORAS_NECESARIAS': (PER, r'^horas de servicio a cubrir al a[ñn]o', 2, 'int'),
    'HORAS_CONTRATADAS': (PER, r'^horas contratadas al a[ñn]o', 2, 'int'),
    # Escenarios (B = pesimista, D = optimista)
    'ESC_PES_INVITADOS': (ESC, r'^invitados por evento$', 2, 'int'),
    'ESC_OPT_INVITADOS': (ESC, r'^invitados por evento$', 4, 'int'),
    'ESC_PES_PRECIO': (ESC, r'^precio medio por invitado sin iva', 2, 'eur2'),
    'ESC_OPT_PRECIO': (ESC, r'^precio medio por invitado sin iva', 4, 'eur2'),
    'ESC_PES_EVENTOS': (ESC, r'^eventos al a[ñn]o', 2, 'int'),
    'ESC_OPT_EVENTOS': (ESC, r'^eventos al a[ñn]o', 4, 'int'),
    'ESC_PES_VENTAS': (ESC, r'^ingresos anuales', 2, 'eur'),
    'ESC_OPT_VENTAS': (ESC, r'^ingresos anuales', 4, 'eur'),
    'ESC_PES_NETO': (ESC, r'^resultado neto$', 2, 'eur'),
    'ESC_OPT_NETO': (ESC, r'^resultado neto$', 4, 'eur'),
}


def fmt(v, f):
    return _fmt(v, f).replace('-', '−', 1) if v < 0 else _fmt(v, f)   # signo menos tipográfico


def _fmt(v, f):
    if f in ('int', 'int_up'):
        n = math.ceil(v - 1e-9) if f == 'int_up' else round(v)
        return '{:,}'.format(int(n)).replace(',', '.')
    if f == 'eur':
        return '{:,}'.format(int(round(v))).replace(',', '.') + ' €'
    if f == 'eur2':
        return '{:,.2f}'.format(v).replace(',', 'X').replace('.', ',').replace('X', '.') + ' €'
    if f in ('pct', 'pct0'):
        txt = ('{:.1f}' if f == 'pct' else '{:.0f}').format(v * 100).replace('.', ',')
        return txt + ' %'
    if f in ('dec1', 'dec2'):
        return ('{:.1f}' if f == 'dec1' else '{:.2f}').format(v).replace('.', ',')
    raise ValueError(f)


def _fila(ws, rx):
    for r in range(1, ws.max_row + 1):
        a = ws.cell(r, 1).value
        if isinstance(a, str) and re.search(rx, a.strip(), re.I):
            return r
    return None


def leer(wb):
    crudo, falta, textos = {}, [], {}
    for tok, (hoja, rx, col, f) in MAPA.items():
        ws = wb[hoja]
        r = _fila(ws, rx)
        v = ws.cell(r, col).value if r else None
        if isinstance(v, str) and tok.startswith('PAYBACK') and v.strip():   # «Más de 3 años»: texto del libro
            t = v.strip()
            textos[tok] = t[:1].lower() + t[1:]   # revisión final m1: va en mitad de frase («se recupera en más de 3 años»)
            continue
        if not isinstance(v, (int, float)):
            falta.append('%s (%s · %s): %r' % (tok, hoja, rx, v))
            continue
        crudo[tok] = (float(v), f)
    c = dict((k, x[0]) for k, x in crudo.items())
    # derivados
    if all(k in c for k in ('CAPEX', 'FIANZA', 'RENTA_PREVIA')):       # las 19 partidas + imprevistos (SPEC D7)
        crudo['INVERSION_PARTIDAS'] = (c['CAPEX'] - c['FIANZA'] - c['RENTA_PREVIA'], 'eur')
    if all(k in c for k in ('VENTAS_A1', 'BE_INGRESOS')):               # holgura sobre el equilibrio CONTABLE, año 1
        crudo['HOLGURA_CONTABLE'] = (c['VENTAS_A1'] / c['BE_INGRESOS'] - 1, 'pct')
    if 'MES_SALDO_MINIMO' in c:
        textos['MES_SALDO_MINIMO_NOMBRE'] = MESES[int(round(c['MES_SALDO_MINIMO'])) - 1]
    if all(k in c for k in ('ORIGEN_FONDOS', 'NECESIDAD_CAJA')):
        crudo['EXCESO_FINANCIACION'] = (c['ORIGEN_FONDOS'] - c['NECESIDAD_CAJA'], 'eur')
    # eventos al año de equilibrio con el tamaño medio de evento del año 1 (camino §3: BE_EVENTOS derivados)
    if all(k in c for k in ('BE_INGRESOS', 'BE_INGRESOS_CAJA', 'PRECIO_INVITADO', 'INVITADOS_EVENTO')):
        evento = c['PRECIO_INVITADO'] * c['INVITADOS_EVENTO']
        crudo['VENTA_EVENTO_MEDIO'] = (evento, 'eur')
        crudo['BE_EVENTOS_ANO'] = (c['BE_INGRESOS'] / evento, 'int_up')
        crudo['BE_EVENTOS_ANO_CAJA'] = (c['BE_INGRESOS_CAJA'] / evento, 'int_up')
        for a in ('A2', 'A3'):
            if 'BE_INGRESOS_CAJA_' + a in c:
                crudo['BE_EVENTOS_CAJA_' + a] = (c['BE_INGRESOS_CAJA_' + a] / evento, 'int_up')
    # el crecimiento lo aplica el motor a los invitados por evento: en euros = eventos del mismo tamaño (SPEC D3)
    for a in ('A2', 'A3'):
        if all(k in c for k in ('VENTAS_' + a, 'VENTAS_A1', 'EVENTOS_ANO')):
            crudo['EVENTOS_EQUIVALENTES_' + a] = (c['EVENTOS_ANO'] * c['VENTAS_' + a] / c['VENTAS_A1'], 'int')
    if 'HORAS_CONTRATADAS' in c:     # jornadas completas equivalentes (40 h × 46 semanas): la columna C de Personal
        crudo['JORNADAS_EQUIVALENTES'] = (c['HORAS_CONTRATADAS'] / (40 * 46.0), 'dec2')   # suma % sin las personas
    if all(k in c for k in ('CANON_A1', 'VENTAS_A1')):          # cada punto de canon (supuesto del 12 %), en ventas
        crudo['CANON_PUNTO'] = (c['CANON_A1'] / 12.0, 'eur')
    return crudo, falta, textos


def prueba_d4(wb):
    """SPEC D4: las dos líneas al 21 % en el P&L, el IVA repercutido de la Tesorería = 21 % de las ventas de cada mes, y
    cada liquidación trimestral (303) = repercutido − soportado del trimestre − lo arrastrado. Devuelve (ok, líneas)."""
    out, ok = [], True
    pyg, tes = wb[PYG], wb[TES]
    tipos = []
    for rx in (r'^ventas de alquiler del espacio', r'^canon del catering exclusivo'):
        r = _fila(pyg, rx)
        g = pyg.cell(r, 7).value
        tipos.append(g)
        out.append('P&L!G%d «%s» = %r' % (r, pyg.cell(r, 1).value, g))
        ok &= isinstance(g, (int, float)) and abs(g - 0.21) < 1e-9
    ventas = pyg.cell(_fila(pyg, r'^ingresos totales'), 2).value
    r_rep, r_sop = _fila(tes, r'^iva repercutido del mes'), _fila(tes, r'^iva soportado del mes')
    r_rep_peso, r_arr = _fila(tes, r'^reparto de la actividad'), _fila(tes, r'^iva a compensar arrastrado')
    r_liq, r_303 = _fila(tes, r'^resultado de la liquidaci'), _fila(tes, r'^pago del iva \(modelo 303\)')
    rep = [tes.cell(r_rep, 2 + i).value for i in range(12)]
    sop = [tes.cell(r_sop, 2 + i).value for i in range(12)]
    peso = [tes.cell(r_rep_peso, 2 + i).value for i in range(12)]
    dif_mes = max(abs(rep[i] - 0.21 * ventas * peso[i]) for i in range(12))
    tot = tes.cell(r_rep, 14).value
    out.append('Tesorería fila %d (IVA repercutido): Σ12 meses = %.2f € = %.4f × ventas del año 1 (%.2f €); máx. '
               'desvío mes a mes frente a 21 %% × ventas del mes = %.4f €' % (r_rep, tot, tot / ventas, ventas, dif_mes))
    ok &= dif_mes < 0.01 and abs(tot - 0.21 * ventas) < 0.01
    for col in range(2, 14):
        liq = tes.cell(r_liq, col).value
        if not isinstance(liq, (int, float)):
            continue
        i = col - 2                         # la liquidación del Mes n cubre los tres meses anteriores
        trimestre = range(i - 3, i)
        arr = tes.cell(r_arr, col).value or 0
        esperado = sum(rep[j] for j in trimestre) - sum(sop[j] for j in trimestre) - arr
        out.append('303 del Mes %d: repercutido %.2f − soportado %.2f − arrastrado %.2f = %.2f (libro %.2f)' % (
            i + 1, sum(rep[j] for j in trimestre), sum(sop[j] for j in trimestre), arr, esperado, liq))
        ok &= abs(esperado - liq) < 0.01
    out.append('Pago del 303 en el año 1 (fila %d, total): %r €' % (r_303, tes.cell(r_303, 14).value))
    return ok, out


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--xlsx', required=True)
    ap.add_argument('--salida', default=SALIDA)
    a = ap.parse_args()
    wb = openpyxl.load_workbook(a.xlsx, data_only=True)
    crudo, falta, textos = leer(wb)
    ok_d4, lineas_d4 = prueba_d4(wb)
    tokens = dict((k, fmt(v, f)) for k, (v, f) in crudo.items())
    tokens.update(textos)
    json.dump({'fuente': os.path.basename(a.xlsx), 'tokens': tokens,
               'valores': dict((k, v) for k, (v, _f) in crudo.items())},
              open(a.salida, 'w', encoding='utf-8'), ensure_ascii=False, indent=1, sort_keys=True)
    v = dict((k, x[0]) for k, x in crudo.items())
    gates = [
        ('DSCR mínimo ≥ 1,30', v.get('DSCR_MIN', 0) >= 1.30),
        ('saldo mínimo de tesorería > 0', v.get('SALDO_MINIMO', -1) > 0),
        ('cobertura de horas ≥ 98 %', v.get('COBERTURA_HORAS', 0) >= 0.98),
        ('holgura sobre el equilibrio de caja del año 2 ≥ 12 %', v.get('HOLGURA_CAJA_A2', -1) >= 0.12),
        ('holgura sobre el equilibrio de caja del año 3 ≥ 15 %', v.get('HOLGURA_CAJA_A3', -1) >= 0.15),
        ('resultado neto del año 3 entre el 6 y el 14 % de las ventas', 0.06 <= v.get('NETO_PCT_A3', -1) <= 0.14),
        ('exceso de financiación (origen - usos) ≥ 0 y < 5 %', 0 <= v.get('DIFERENCIA_FONDOS_PCT', -1) < 0.05),
        ('D4: las dos líneas y el IVA repercutido de la Tesorería al 21 %; 303 = repercutido − soportado − arrastrado',
         ok_d4),
    ]
    for tok in ('INVITADOS_EVENTO', 'PRECIO_INVITADO', 'EVENTOS_ANO', 'VENTAS_A1', 'VENTAS_A2', 'VENTAS_A3',
                'EVENTOS_EQUIVALENTES_A2', 'EVENTOS_EQUIVALENTES_A3', 'RAI_A1', 'NETO_A1', 'NETO_A2', 'NETO_A3',
                'NETO_PCT_A1', 'NETO_PCT_A2', 'NETO_PCT_A3', 'INVERSION_PARTIDAS', 'CAPEX', 'INVERSION_TOTAL',
                'NECESIDAD_CAJA', 'RECURSOS_PROPIOS', 'PRESTAMO', 'DIFERENCIA_FONDOS', 'DIFERENCIA_FONDOS_PCT',
                'CUOTA_ANUAL', 'DSCR_A2', 'DSCR_A3', 'DSCR_MIN', 'SALDO_MINIMO', 'MES_SALDO_MINIMO',
                'MES_SALDO_MINIMO_NOMBRE', 'SALDO_MES12', 'BE_INVITADOS_EVENTO', 'BE_INVITADOS_EVENTO_CAJA',
                'BE_EVENTOS_ANO', 'BE_EVENTOS_ANO_CAJA', 'BE_EVENTOS_CAJA_A2', 'BE_EVENTOS_CAJA_A3',
                'HOLGURA_CONTABLE', 'HOLGURA_CAJA', 'HOLGURA_CAJA_A2', 'HOLGURA_CAJA_A3', 'COBERTURA_HORAS', 'PAYBACK',
                'MARGEN_BRUTO_PCT_A1', 'PERSONAL_PCT_A1', 'RENTA_PCT_A1', 'ESC_PES_NETO', 'ESC_OPT_NETO'):
        print('  %-26s %s' % (tok, tokens.get(tok, '—')))
    print('  D4 (lectura del xlsx calculado):')
    for x in lineas_d4:
        print('    ' + x)
    for nombre, ok in gates:
        print('  [%s] %s' % ('OK' if ok else 'ROJO', nombre))
    print('  [info] holgura de caja del año 1: %s · neto del año 1: %s' % (
        tokens.get('HOLGURA_CAJA', '—'), tokens.get('NETO_PCT_A1', '—')))
    for f in falta:
        print('  FALTA', f)
    ok_total = all(ok for _n, ok in gates) and not falta
    print('cifras_caso_es.json: %d tokens · gates del caso D9 %s' % (len(tokens), 'VERDE' if ok_total else 'ROJA'))
    sys.exit(0 if ok_total else 1)


if __name__ == '__main__':
    main()
