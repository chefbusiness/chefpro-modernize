#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""validar_mercado.py — gate de mercado_en.json (Recipe Costing Kit Pro, F2-EN).

SPEC: D10-D12 y §4 gate 4 («Ejemplos creíbles»). Sesión Claude Code, 24-sep-2026.
Lee los xlsx ES v2.1 publicados SOLO para leer (nunca los escribe) y recalcula todo con los factores de
mapas.conversions(), sin fiarse de lo que calculó generar_mercado.py.

Comprueba
  A. Fuentes: todo precio tiene fuente (BLS / USDA / distribuidor / secundaria, con URL y fecha) o
     [estimado]/[derivado] con razonamiento (D11). Unidades en la lista D7, categorías EN válidas.
  B. Cobertura: toda fila de ingrediente de las fichas ES 01-08 tiene su fila EN (y ninguna sobra); los 15
     productos del BONUS; las filas de «Formatos de Compra» del 04 (+ el barril de D7 = 8); el 13 = los
     ingredientes distintos de las recetas EN.
  C. Filas: precio = el del catálogo; la clave «compra→uso» existe en Conversions (la ficha la resuelve).
  D. Cantidades: conversión honesta desde el ES (±25 %) o cambio documentado.
  E. Food cost implícito (coste por ración ÷ precio de carta de referencia) dentro de D12; pour cost del 04
     entre 18 y 24 %; menús 02 y 03 a nivel de menú.
  F. Libro 12: merma recalculada = la de la fila 5 de la 01; coste por porción igual (0,5 %); 0 < rendimiento
     < 1; factor ≥ 1. Cocción: coste = ficha 08.
  G. Libro 13: formato ÷ contenido = precio de la ficha ÷ factor (0,5 %); unidades base en D7; la alerta
     salta con la demo del aceite y no con la del solomillo.
  H. BONUS: la desviación relativa H/F de cada producto = la del ES.
  I. Valores por defecto (D6 impuesto 0, Q-factor 10 %, 06, 08, 10) y textos (sin €, sin no latinos, rangos
     D12 junto a su tipo de local).

Uso:  python3 validar_mercado.py [--verbose] [--json RUTA]   → sale con código 1 si algo falla.
"""
from __future__ import print_function

import json
import math
import os
import re
import sys
from collections import OrderedDict, Counter

AQUI = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, AQUI)
import mapas  # noqa: E402

RAIZ = os.path.abspath(os.path.join(AQUI, '..', '..', '..'))
ES_DIR = os.path.join(RAIZ, 'astro-site', 'public', 'dl', 'kit-escandallos')
_JSON = sys.argv[sys.argv.index('--json') + 1] if '--json' in sys.argv else os.path.join(AQUI, 'mercado_en.json')
M = json.load(open(_JSON, encoding='utf-8'))
FAC = {r['clave']: r['factor'] for r in mapas.conversions()}
D7 = set(mapas.UNIDADES_EN)
CATS_EN = set(mapas.CATEGORIAS.values())
VERBOSE = '--verbose' in sys.argv
FALLOS = []
AVISOS = []
TOL = 0.005


def fallo(msg):
    FALLOS.append(msg)


def aviso(msg):
    AVISOS.append(msg)


def fac(a, b):
    return FAC.get(a + '→' + b)


# ---------------------------------------------------------------- A. fuentes
TIPOS = {'BLS', 'USDA AMS', 'distribuidor', 'fuente secundaria (resumen de buscador)', '[estimado]', '[derivado]'}


def fuente_ok(f, donde):
    if not isinstance(f, dict) or f.get('tipo') not in TIPOS:
        fallo('%s: sin fuente válida (%r)' % (donde, f if not isinstance(f, dict) else f.get('tipo')))
        return None
    t = f['tipo']
    if t == '[estimado]':
        if len(f.get('razonamiento', '')) < 30:
            fallo('%s: [estimado] sin razonamiento' % donde)
    elif t == '[derivado]':
        if len(f.get('razonamiento', '')) < 20:
            fallo('%s: [derivado] sin razonamiento' % donde)
        for p in f.get('partes', []):
            fuente_ok(p, donde + ' (parte)')
    else:
        if not str(f.get('url', '')).startswith('http') or not f.get('fecha'):
            fallo('%s: fuente %s sin URL o fecha' % (donde, t))
    return t


cat = M['catalogo']
tipos_cat = Counter()
for iid, c in cat.items():
    t = fuente_ok(c['fuente'], 'catálogo ' + iid)
    tipos_cat[t] += 1
    if not (isinstance(c['precio'], (int, float)) and c['precio'] > 0):
        fallo('catálogo %s: precio no positivo' % iid)
    if c['ud_compra'] not in D7:
        fallo('catálogo %s: unidad de compra %r fuera de D7' % (iid, c['ud_compra']))
    for u in c.get('otras_ud_compra', {}):
        if u not in D7:
            fallo('catálogo %s: unidad %r fuera de D7' % (iid, u))
    if c['categoria_en'] not in CATS_EN:
        fallo('catálogo %s: categoría %r' % (iid, c['categoria_en']))
    if t == 'distribuidor' and c['fuente'].get('distribuidor') == 'WebstaurantStore':
        # el precio de la ficha sale del listado: precio del formato ÷ contenido en su unidad
        f = c['fuente']
        u = f['unidad_contenido']
        base = f['precio_formato'] / f['contenido_formato']
        esperado = base if u == c['ud_compra'] else base * (fac(c['ud_compra'], u) or float('nan'))
        if not abs(esperado - c['precio']) <= max(0.005, esperado * TOL):
            fallo('catálogo %s: precio %.4f ≠ listado %.4f' % (iid, c['precio'], esperado))


# ---------------------------------------------------------------- B. cobertura (ES publicado)
def leer_fichas_es():
    import openpyxl
    libros = ['01-escandallo-estandar.xlsx', '02-menu-degustacion.xlsx', '03-menu-del-dia.xlsx',
              '04-cocktails-bebidas.xlsx', '05-pasteleria.xlsx', '06-catering.xlsx', '07-cafeteria-brunch.xlsx',
              '08-food-truck.xlsx']
    filas = OrderedDict()
    for f in libros:
        wb = openpyxl.load_workbook(os.path.join(ES_DIR, f), read_only=False)
        for ws in wb.worksheets:
            hdr = None
            for r in range(1, 10):
                if ws.cell(r, 1).value == 'Ingrediente':
                    hdr = r
                    break
            if not hdr:
                continue
            r = hdr + 1
            while r < 80:
                a = ws.cell(r, 1).value
                if isinstance(a, str) and a.upper().startswith('COSTE TOTAL'):
                    break
                if a is not None:
                    filas[(f[:2], ws.title, r)] = OrderedDict([
                        ('nombre', a), ('ud_c', ws.cell(r, 3).value), ('cant', ws.cell(r, 5).value),
                        ('ud_u', ws.cell(r, 6).value), ('merma', ws.cell(r, 8).value)])
                r += 1
    return filas


ES = leer_fichas_es()
EN = OrderedDict()
for rec in M['recetas']:
    for f in rec['filas']:
        EN[(rec['libro'], rec['hoja_es'], f['fila'])] = (rec, f)
for k, v in ES.items():
    if k not in EN:
        fallo('fila ES sin fila EN: %s %s fila %d «%s»' % (k[0], k[1], k[2], v['nombre']))
    elif EN[k][1]['nombre_es'] != v['nombre']:
        fallo('fila %s %s %d: nombre ES %r ≠ %r' % (k[0], k[1], k[2], EN[k][1]['nombre_es'], v['nombre']))
for k in EN:
    if k not in ES:
        fallo('fila EN sin fila ES: %s %s fila %d' % k)


def leer_otros_es():
    import openpyxl
    wb = openpyxl.load_workbook(os.path.join(ES_DIR, 'BONUS-mermas-inventario.xlsx'))
    ws = wb['Inventario']
    bonus = [(r, ws.cell(r, 1).value, ws.cell(r, 2).value, ws.cell(r, 3).value, ws.cell(r, 4).value,
              ws.cell(r, 5).value) for r in range(5, 20)]
    vs = wb['Ventas del periodo']
    platos = []
    for r in range(5, 15):
        if vs.cell(r, 1).value:
            q = OrderedDict()
            for c in range(3, 18):
                if vs.cell(r, c).value:
                    q[ws.cell(5 + c - 3, 1).value] = vs.cell(r, c).value
            platos.append((vs.cell(r, 1).value, vs.cell(r, 2).value, q))
    wb4 = openpyxl.load_workbook(os.path.join(ES_DIR, '04-cocktails-bebidas.xlsx'))
    fc = wb4['Formatos de Compra']
    bot = [(r, fc.cell(r, 1).value) for r in range(5, 17) if fc.cell(r, 1).value]
    return bonus, platos, bot


BONUS_ES, PLATOS_ES, BOT_ES = leer_otros_es()
prods = {p['fila']: p for p in M['bonus']['productos']}
for (r, nom, ud, c, d, e) in BONUS_ES:
    p = prods.get(r)
    if not p or p['producto_es'] != nom:
        fallo('BONUS fila %d «%s» sin producto EN' % (r, nom))
        continue
    if p['id'] not in cat:
        fallo('BONUS %s: id %s no está en el catálogo' % (nom, p['id']))
if len(M['bottle_sizes']) != len(BOT_ES) + 1:
    fallo('Bottle Sizes: %d filas EN para %d ES + 1 barril' % (len(M['bottle_sizes']), len(BOT_ES)))
for b in M['bottle_sizes']:
    fuente_ok(b['fuente'], 'Bottle Sizes ' + b['producto'])
    if abs(b['precio_litro'] - round(b['precio_botella'] * 1000 / b['ml'], 4)) > 1e-4:
        fallo('Bottle Sizes %s: precio por litro ≠ botella × 1000 ÷ ml (E3)' % b['producto'])

# ---------------------------------------------------------------- C/D/E. filas, cantidades, food cost
BASE = {'kg': ('g', 1000), 'g': ('g', 1), 'L': ('ml', 1000), 'ml': ('ml', 1), 'cl': ('ml', 10),
        'ud': ('each', 1), 'docena': ('each', 12), 'manojo': ('bunch', 1)}


def en_base(q, u):
    if u in ('oz', 'lb', 'kg', 'g'):
        return 'g', q * fac(u, 'g')
    if u in ('fl oz', 'tsp', 'tbsp', 'cup', 'qt', 'pt', 'gal', 'L', 'ml', 'cl'):
        return 'ml', q * fac(u, 'ml')
    if u in ('each', 'dozen'):
        return 'each', q * fac(u, 'each')
    if u == 'bunch':
        return 'bunch', q
    return None, None


tabla = []
for rec in M['recetas']:
    tot = 0.0
    for f in rec['filas']:
        donde = '%s %s fila %d' % (rec['libro'], rec['hoja_en'], f['fila'])
        c = cat.get(f['id'])
        if not c:
            fallo(donde + ': id sin catálogo')
            continue
        precio = c['precio'] if f['ud_compra'] == c['ud_compra'] else c.get('otras_ud_compra', {}).get(f['ud_compra'])
        if precio is None or abs(precio - f['precio']) > 1e-9:
            fallo(donde + ': precio de la fila ≠ catálogo')
        for u in (f['ud_compra'], f['ud_uso']):
            if u not in D7:
                fallo(donde + ': unidad %r fuera de D7' % u)
        g = fac(f['ud_compra'], f['ud_uso'])
        if g is None:
            fallo(donde + ': la clave «%s→%s» no está en Conversions (la ficha diría «check units»)'
                  % (f['ud_compra'], f['ud_uso']))
            continue
        if not (0 <= f['merma'] < 1):
            fallo(donde + ': merma fuera de [0,1)')
        if f['categoria_en'] not in CATS_EN:
            fallo(donde + ': categoría EN %r' % f['categoria_en'])
        coste = f['cantidad'] / (1 - f['merma']) * f['precio'] / g
        tot += coste
        # D. conversión honesta desde el ES
        es = ES.get((rec['libro'], rec['hoja_es'], f['fila']))
        if es:
            be, qe = BASE.get(es['ud_u'], (None, None))
            bn, qn = en_base(f['cantidad'], f['ud_uso'])
            if be and bn and be == bn and qe:
                ratio = qn / (es['cant'] * qe)
                if abs(ratio - 1) > 0.25 and not f.get('cambio_cantidad') and not f.get('nota'):
                    fallo(donde + ': cantidad EN %.3g %s = %.0f%% del ES sin explicar' % (f['cantidad'], f['ud_uso'], ratio * 100))
                elif abs(ratio - 1) > 0.25:
                    aviso(donde + ': cantidad ×%.2f del ES (documentado)' % ratio)
            elif be != bn and not f.get('nota') and not f.get('cambio_cantidad'):
                fallo(donde + ': cambia de dimensión (%s → %s) sin nota' % (be, bn))
    coste_racion = tot * (1 + rec['q_factor']) / rec['raciones']
    if abs(coste_racion - rec['coste_racion']) > 1e-4:
        fallo('%s %s: coste por ración recalculado %.4f ≠ %.4f' % (rec['libro'], rec['hoja_en'], coste_racion, rec['coste_racion']))
    rec['_coste'] = coste_racion
    if rec['precio_por'] == 'receta':
        fuente_ok(rec['precio_carta_fuente'], 'precio de carta %s %s' % (rec['libro'], rec['hoja_en']))
        fc = coste_racion / rec['precio_carta_ref']
        lo, hi = M['d12'][rec['tipo_local']]['min'], M['d12'][rec['tipo_local']]['max']
        ok = lo <= fc <= hi
        tabla.append((rec['libro'], rec['hoja_en'], coste_racion, rec['precio_carta_ref'], fc, lo, hi, ok))
        if not ok:
            fallo('%s %s: food cost implícito %.1f%% fuera de D12 %s (%d-%d%%)'
                  % (rec['libro'], rec['hoja_en'], fc * 100, rec['tipo_local'], lo * 100, hi * 100))
    if rec['libro'] == '04' and not (0.18 <= coste_racion / rec['precio_carta_ref'] <= 0.24):
        fallo('04 %s: pour cost fuera de 18-24 %%' % rec['hoja_en'])
    if rec['q_factor'] != (0.05 if rec['libro'] == '04' else 0.10):
        fallo('%s %s: Q-factor %.2f' % (rec['libro'], rec['hoja_en'], rec['q_factor']))

for lib, key, extra in (('02', 'menu_02', 0.0), ('03', 'menu_03', None)):
    meta = M[key]
    fuente_ok(meta['fuente'], 'precio del menú ' + lib)
    if extra is None:
        ex = meta['extras']
        fuente_ok(ex['bread_and_butter']['fuente'], 'menú 03 bread & butter')
        extra = ex['bread_and_butter']['valor'] + ex['drink']['valor'] + ex['coffee']['valor']
        if ex['drink']['valor'] != 0 or ex['coffee']['valor'] != 0:
            fallo('menú 03: bebida y café deben ir a 0 por defecto (M19)')
    total = sum(r['_coste'] for r in M['recetas'] if r['libro'] == lib) + extra
    fc = total / meta['precio_carta_ref']
    lo, hi = M['d12'][meta['tipo_local']]['min'], M['d12'][meta['tipo_local']]['max']
    ok = lo <= fc <= hi
    tabla.append((lib, 'MENU (%s pases)' % sum(1 for r in M['recetas'] if r['libro'] == lib), total,
                  meta['precio_carta_ref'], fc, lo, hi, ok))
    if not ok:
        fallo('menú %s: food cost %.1f%% fuera de D12' % (lib, fc * 100))
    for r in M['recetas']:
        # CHEF-03 (opción A): los pases no se venden sueltos → «Current menu price» vacío, como en el ES
        if r['libro'] == lib and r['precio_carta_ref'] is not None:
            fallo('menú %s: el pase %s lleva precio propio (%r); debe ir vacío' % (lib, r['hoja_en'], r['precio_carta_ref']))
    nota = meta.get('nota_resumen') or {}
    t = nota.get('texto', '')
    if '${:,.0f}'.format(meta['precio_carta_ref']) not in t or '{:.1f}%'.format(fc * 100) not in t:
        fallo('menú %s: la nota del Summary no cita el precio del menú y su food cost (%r)' % (lib, t[:80]))

# ---------------------------------------------------------------- F. libro 12
d = M['libro12']['despiece']
comp = d['componentes']
ap = d['C8_peso_ap_lb']
p = d['C9_precio_lb']
util = sum(c['peso_lb'] for c in comp if c['tipo'] == 'Usable')
subp = sum(c['peso_lb'] * c['valor_lb'] for c in comp if c['tipo'] == 'By-product')
corte = ap - sum(c['peso_lb'] for c in comp)
coste_lb = (ap * p - subp) / util
merma = 1 - p / coste_lb
if abs(corte) > 1e-4:
    fallo('libro 12: la pérdida de corte no es 0 (%.4f lb)' % corte)
if not (0 < util / ap < 1) or coste_lb / p < 1:
    fallo('libro 12: rendimiento o factor de coste fuera de rango')
for c in comp:
    if len(c.get('fuente', '')) < 20:
        fallo('libro 12: componente «%s» sin fuente' % c['pieza'])
if set(c['tipo'] for c in comp) - set(mapas.DV12.values()):
    fallo('libro 12: tipo fuera de la DV EN')
r01 = ([f for r in M['recetas'] if r['libro'] == '01' for f in r['filas'] if f['fila'] == 5] or [None])[0]
if r01 is None:
    print('✗ la 01 EN no tiene fila 5 (solomillo): no se puede comprobar el libro 12'); sys.exit(1)
if abs(round(merma, 4) - r01['merma']) > 1e-9:
    fallo('libro 12: merma %.4f ≠ fila 5 de la 01 (%.4f)' % (merma, r01['merma']))
if abs(r01['precio'] - p) > 1e-9:
    fallo('libro 12: precio AP ≠ el de la 01')
c12 = d['C36_porcion_oz'] / d['C11_unidades_porcion'] * coste_lb
c01 = r01['cantidad'] / (1 - r01['merma']) * r01['precio'] / fac(r01['ud_compra'], r01['ud_uso'])
if abs(c12 - c01) / c12 > TOL:
    fallo('libro 12: coste por porción %.4f ≠ ficha 01 %.4f' % (c12, c01))
if d['C11_unidades_porcion'] != 16 or r01['cantidad'] != d['C36_porcion_oz']:
    fallo('libro 12: celda de porción ≠ 16 o porción ≠ la de la 01')
k = M['libro12']['coccion']
cc = k['C12_porcion_cocinada_oz'] / 16.0 * k['C8_precio_lb'] / ((1 - k['C10_merma_despiece']) * (k['C9_peso_cocinado_lb'] / k['C7_peso_crudo_lb']))
r08 = [f for r in M['recetas'] if r['libro'] == '08' and r['hoja_es'] == 'Smash Burger' for f in r['filas'] if f['fila'] == 5][0]
c08 = r08['cantidad'] / (1 - r08['merma']) * r08['precio'] / fac(r08['ud_compra'], r08['ud_uso'])
if abs(cc - c08) / c08 > TOL or abs(k['C8_precio_lb'] - r08['precio']) > 1e-9:
    fallo('libro 12 cocción: coste %.4f ≠ ficha 08 %.4f' % (cc, c08))

# ---------------------------------------------------------------- G. libro 13
l13 = M['libro13']
ids13 = set(f['id'] for f in l13['filas'])
ids_rec = set(f['id'] for r in M['recetas'] for f in r['filas'] if cat[f['id']].get('formato13'))
if ids13 != ids_rec:
    fallo('libro 13: ingredientes ≠ los de las recetas EN (faltan %s, sobran %s)' % (sorted(ids_rec - ids13), sorted(ids13 - ids_rec)))
if l13['n_ingredientes'] != len(l13['filas']):
    fallo('libro 13: n_ingredientes no cuadra')
for f13 in l13['filas']:
    if f13['unidad_base'] not in D7:
        fallo('libro 13 %s: unidad base %r fuera de D7' % (f13['ingrediente'], f13['unidad_base']))
    pub = f13['precio_formato'] / f13['contenido']
    for r in M['recetas']:
        for f in r['filas']:
            if f['id'] != f13['id']:
                continue
            g = fac(f['ud_compra'], f13['unidad_base'])
            if g is None:
                fallo('libro 13 %s: sin factor %s→%s' % (f13['ingrediente'], f['ud_compra'], f13['unidad_base']))
                continue
            if abs(pub - f['precio'] / g) / pub > TOL:
                fallo('libro 13 %s: %.4f por %s ≠ ficha %s %s %.4f' % (f13['ingrediente'], pub, f13['unidad_base'],
                                                                      r['libro'], r['hoja_en'], f['precio'] / g))
    if 'precio_anterior' in f13:
        var = round(f13['precio_unidad_base'] / f13['precio_anterior'] - 1, 6)
        alerta = var > l13['umbral_alerta']
        if (f13['id'] == 'evoo') != alerta:
            fallo('libro 13 %s: la demo de alerta no hace lo que dice (%.2f%%)' % (f13['ingrediente'], var * 100))
if sum(1 for f in l13['filas'] if 'precio_anterior' in f) != 2:
    fallo('libro 13: tiene que haber exactamente 2 filas de demostración (R2-10)')

# ---------------------------------------------------------------- H. BONUS
def dev_rel(stock, platos, nombre):
    c, dd, e = stock
    F = c + dd - e
    G = sum(rac * q.get(nombre, 0) for (rac, q) in platos)
    return (F - G) / F if F else None


pl_es = [(rac, q) for (_n, rac, q) in PLATOS_ES]
pl_en = [(pl['raciones'], {k2: v for k2, v in pl['por_racion'].items()}) for pl in M['bonus']['platos']]
for (r, nom, ud, c, dd, e) in BONUS_ES:
    p = prods.get(r)
    if c is None or p is None:
        continue
    es_dev = dev_rel((c, dd, e), pl_es, nom)
    en_dev = dev_rel((p['stock_ini'], p['compras'], p['stock_fin']), pl_en, nom)
    # M10 a ±0,0005 (0,05 puntos): las compras van en cajas enteras y el stock en 1-2 decimales (CHEF-13).
    if es_dev is None or en_dev is None or abs(es_dev - en_dev) > 5e-4:
        fallo('BONUS %s: desviación relativa EN %s ≠ ES %s' % (nom, en_dev, es_dev))
anc = M['bonus']['anclas_ap_01']
f01 = {f['fila']: f for r in M['recetas'] if r['libro'] == '01' for f in r['filas']}
for nom, fila in (('Solomillo ternera', 5), ('Patata', 6), ('Mantequilla', 8)):
    if fila not in f01:
        fallo('BONUS %s: la 01 EN no tiene fila %d para anclar la ración' % (nom, fila))
        continue
    ap_lb = f01[fila]['cantidad'] / (1 - f01[fila]['merma']) / 16.0
    q = M['bonus']['platos'][0]['por_racion'][nom]
    if abs(q - ap_lb) > 5e-4:
        fallo('BONUS %s: ración %.4f lb ≠ AP qty de la 01 EN %.4f lb' % (nom, q, ap_lb))
for p in M['bonus']['productos']:
    c = cat[p['id']]
    esperado = c['precio'] if p['ud_en'] == c['ud_compra'] else c['precio'] / fac(c['ud_compra'], p['ud_en'])
    if abs(esperado - p['precio_unitario']) > max(0.005, esperado * TOL):
        fallo('BONUS %s: precio %.4f ≠ catálogo %.4f' % (p['producto_en'], p['precio_unitario'], esperado))
    if p['ud_en'] not in D7:
        fallo('BONUS %s: unidad %s fuera de D7' % (p['producto_en'], p['ud_en']))

# ---------------------------------------------------------------- I. valores por defecto y textos
ev = M['catering_06']
esper06 = {'C9_invitados_por_camarero': 25, 'C12_coste_hora_camarero': 25.0, 'C14_coste_hora_jefe_sala': 35.0,
           'C17_rentals_por_invitado': 3.5, 'C28_minimo': 1000.0, 'C25_markup_servicios': 0.20, 'C31_impuesto': 0.0}
for kk, vv in esper06.items():
    if ev[kk] != vv:
        fallo('06 %s = %r (SPEC §2.5: %r)' % (kk, ev[kk], vv))
if ev['C25_rotulo'] != 'Markup on services (%)':
    fallo('06 C25: rótulo')
be = M['break_even_08']
if [l['valor'] for l in be['lineas']] != [45.0, 25.0, 30.0, 300.0, 50.0, 15.0] or be['total'] != 465.0 or be['B25_impuesto'] != 0:
    fallo('08 break-even ≠ 45+25+30+300+50+15 = 465 $/día con impuesto 0')
c10 = M['calculadora_10']
for fila in c10['filas']:
    d12 = M['d12'][fila['tipo']]
    if (fila['fc_min'], fila['fc_max']) != (d12['min'], d12['max']):
        fallo('10 %s: rango ≠ D12' % fila['rotulo'])
    if fila['comision'] != (0.25 if fila['tipo'] == 'delivery' else 0.0):
        fallo('10 %s: comisión' % fila['rotulo'])
if c10['C5_impuesto'] != 0 or M['dashboard_11']['D4_objetivo'] != 0.30:
    fallo('10/11: impuesto 0 y objetivo 30 %')
for key in ('waste_09', 'dashboard_11'):
    fuente_ok(M[key]['fuente'], key)
fuente_ok(c10['C4_fuente'], 'calculadora 10 C4')
w9 = M['waste_09']
for fam in w9['familias_con_ejemplo']:
    if not (fam['compra'] > 0 and 0 < fam['desperdicio'] < fam['compra']):
        fallo('09 %s: compra/desperdicio' % fam['familia'])

NO_LATINO = re.compile(r'[Ѐ-ӿ֐-׿؀-ۿ฀-๿ᄀ-ᇿ぀-ヿ㐀-鿿가-힯]')
RESTOS_ES = re.compile(r'\b(IVA|€|comensales|escandallo|merma|ración|raciones|plato|precio|según)\b', re.I)
ETIQUETAS = [('pour cost at the bar', 'bar'), ('pastry & bakery', 'pastry_bakery'), ('for catering', 'catering'),
             ('for a café', 'cafe'), ('for a food truck', 'food_truck'), ('Its target food cost', 'delivery')]
for t in M['textos_cifra_derivada']:
    en = t['en']
    if NO_LATINO.search(en):
        fallo('texto con caracteres no latinos: %s' % en[:60])
    if '€' in en or RESTOS_ES.search(en):
        fallo('texto EN con restos de español o €: %s' % en[:80])
    for et, tipo in ETIQUETAS:
        if et in en:
            d12 = M['d12'][tipo]
            rng = '%d-%d%%' % (round(d12['min'] * 100), round(d12['max'] * 100))
            if rng not in en:
                fallo('texto «%s…»: el rango de %s no es el de D12 (%s)' % (en[:50], tipo, rng))
for c in cat.values():
    if NO_LATINO.search(c['nombre_en']) or '€' in c['nombre_en']:
        fallo('nombre EN con caracteres no válidos: %s' % c['nombre_en'])
for r in M['recetas']:
    for f in r['filas']:
        if NO_LATINO.search(f['nombre_en']):
            fallo('fila con caracteres no latinos: %s' % f['nombre_en'])

# Campos que acaban en el xlsx EN: sin restos de español (lista blanca de §4 gate 2: açaí, fleur de sel…).
BLANCA = re.compile(r'(?i)fleur de sel|açaí|acai|ximénez|purée|ibérico|manchego|jalapeño')
ES_PAL = re.compile(r'(?i)\b(de|del|la|las|los|el|y|con|para|por|sin|una|uno|precio|caja|saco|bote|ración|merma|plato|según|también|más)\b')
ES_CAR = re.compile(r'[ñÑáíóúÁÍÓÚ¿¡€]')
campos = []
for c in cat.values():
    campos += [c['nombre_en'], c['proveedor13'] or ''] + ([c['formato13']['texto']] if c['formato13'] else [])
for r in M['recetas']:
    campos += [r['hoja_en'], r['titulo_en'] or ''] + [f['nombre_en'] for f in r['filas']]
campos += [b['producto'] for b in M['bottle_sizes']] + [b['donde'] for b in M['bottle_sizes']]
campos += [M['catering_06']['C25_rotulo']] + [l['rotulo'] for l in M['break_even_08']['lineas']]
campos += [x['rotulo'] for x in M['menu_03']['extras'].values()]
campos += [p['producto_en'] for p in M['bonus']['productos']] + [p['plato_en'] for p in M['bonus']['platos']]
campos += [d['titulo'], d['C5_producto']] + [c['pieza'] for c in comp] + [k['titulo'], k['C5_elaboracion']]
campos += [f['notas'] for f in l13['filas']] + [f['formato'] for f in l13['filas']]
campos += [f['rotulo'] for f in c10['filas']] + [t['en'] for t in M['textos_cifra_derivada']]
for s_ in campos:
    limpio = BLANCA.sub('', s_ or '')
    if ES_PAL.search(limpio) or ES_CAR.search(limpio) or NO_LATINO.search(limpio):
        fallo('campo EN con restos de español: %r' % (s_[:80]))

# Cada texto con cifra derivada tiene que casar con una cadena ES exacta (aplicar_en sustituye por esa clave).
_ES_TXT = set(c['es'] for c in json.load(open(os.path.join(AQUI, 'textos_es.json'), encoding='utf-8'))['cadenas'])
for t in M['textos_cifra_derivada']:
    if t['es'] not in _ES_TXT:
        fallo('texto con cifra derivada sin cadena ES exacta: %r' % t['es'][:60])

# ---------------------------------------------------------------- informe
print('mercado_en.json · validar_mercado.py')
print('  catálogo: %d ingredientes → %s' % (len(cat), dict(tipos_cat)))
todas = [(r, f) for r in M['recetas'] for f in r['filas']]
print('  filas de ficha EN: %d (ES: %d) · BONUS: %d productos · Bottle Sizes: %d · libro 13: %d ingredientes'
      % (len(todas), len(ES), len(M['bonus']['productos']), len(M['bottle_sizes']), len(l13['filas'])))
print('  %-4s %-22s %9s %8s %7s  %s' % ('lib', 'receta', 'coste', 'carta', 'FC', 'D12'))
for (lib, nom, coste, ref, fc, lo, hi, ok) in tabla:
    print('  %-4s %-22s %9.2f %8.2f %6.1f%%  %d-%d%% %s' % (lib, nom[:22], coste, ref, fc * 100, lo * 100, hi * 100,
                                                          'OK' if ok else 'FUERA'))
print('  libro 12: merma %.2f%% = fila 5 de la 01; coste porción %.4f vs %.4f' % (merma * 100, c12, c01))
if VERBOSE:
    for a in AVISOS:
        print('  aviso:', a)
else:
    print('  avisos (cambios de cantidad documentados): %d  (--verbose para verlos)' % len(AVISOS))
if FALLOS:
    print('FALLOS (%d):' % len(FALLOS))
    for f in FALLOS:
        print('  ✗', f)
    sys.exit(1)
print('OK: 0 fallos')
