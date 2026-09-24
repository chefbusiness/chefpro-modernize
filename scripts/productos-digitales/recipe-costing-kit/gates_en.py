#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
gates_en.py — Gates F2 del Recipe Costing Kit Pro (EN), SPEC §4 (1-7bis) sobre la salida de
aplicar_en.py. Solo LEE: los 14 xlsx EN, los 14 xlsx ES v2.1 publicados, censo_es.json,
textos_es.json, mercado_en.json e instrucciones_extra_en.json.

    python3 gates_en.py --dir <scratchpad>/rck-dryrun [--mes September] [--json informe.json]
    python3 gates_en.py --dir … --solo 1,2,5          # solo algunos gates
    python3 gates_en.py --autotest --dir …             # inyecta defectos en una copia y exige que
                                                       # los gates los detecten todos

Gates (SPEC §4 F2):
  1    paridad estructural EN↔ES (hojas por el mapa; fórmula a fórmula salvo mapas de hoja y de
       literales y E1-E5; DV, CF, merges, protección, paneles, áreas y títulos de impresión,
       celdas desbloqueadas, ajuste de página y series de gráfico);
  2    restos: cero «€», «$» solo en texto como importe US, cero no latinos, cero español
       (valores, number_format, styles.xml, nombres de hoja, DV, gráficos, cabeceras/pies, docProps);
  3    cálculo: caché en toda fórmula (inject_cache), cero errores salvo el #N/A buscado de los
       gráficos, cero funciones que pycel no evalúa, los 6 checks de §2.3 y sensibilidad (pycel);
  4    ejemplos creíbles: validar_mercado.py en verde y food cost implícito del xlsx dentro de D12;
  5    formato: Letter, sin «€»/dd/mm, fechas = numFmtId 14 (contadas contra censo_es.json),
       línea de versión D14, pie EN y docProps;
  6    coherencia de texto acotada: rangos D12, cifras derivadas (M5) y pestañas citadas (§2.2);
  7    censo-entregables.py --fail sobre la carpeta, con Letter admitido;
  7bis libros 12 y 13 (rendimiento, factor, coste = ficha 01; alerta; 13 = fichas; unidades D7);
  pdf  el bono PDF contra los xlsx (Letter, nº de páginas, ficheros y pestañas citados, tabla de trim loss =
       «Trim Loss Factors», tabla D12 = calculadora 10, restos). Revisión R1 EN, T11.
Sale con código 1 si algún gate falla.

Revisión R1 EN (24-sep): gate 1 compara en sentido directo (Transformador ES→EN == EN, T1); gate 2 mira también
los literales de fórmula y de CF, busca IDs internos de la SPEC (T2) y usa una LISTA BLANCA de caracteres (T6);
gate 3 exige <v> en el XML de toda celda con <f> (T5) y prueba la celda de tamaño vacía (T7); gate 4 mide los menús
02/03 a nivel de menú con los pases sin precio (CHEF-03); gate 5 exige «Tax rate (%)» = 0 en toda fila (CHEF-01);
gate 6 cita pestañas y rótulos por coincidencia EXACTA (T4). El autotest pasa cada gate por la copia limpia antes
de mutar y exige, por defecto, una subcadena del mensaje esperado (T11).
"""
import argparse
import contextlib
import html
import json
import logging
import os
import re
import shutil
import subprocess
import sys
import zipfile
from collections import OrderedDict, Counter, defaultdict

AQUI = os.path.dirname(os.path.abspath(__file__))
SCRIPTS = os.path.dirname(AQUI)
REPO = os.path.dirname(os.path.dirname(SCRIPTS))
sys.path.insert(0, AQUI)

import openpyxl                                          # noqa: E402

import mapas                                             # noqa: E402
import extraer_textos                                    # noqa: E402
import aplicar_en                                        # noqa: E402  (rutas y constantes)

logging.disable(logging.CRITICAL)

ORIGEN = aplicar_en.ORIGEN
CENSO_SCRIPT = os.path.join(SCRIPTS, 'censo-entregables.py')
VALIDAR_MERCADO = os.path.join(AQUI, 'validar_mercado.py')
INJECT = os.path.join(SCRIPTS, 'inject_cache.py')
TOL = 0.005
PAPEL_LETTER = 1
PIE_EN = 'AI Chef Pro · aichef.pro · Page &P of &N'
PAGINAS_PDF = 23           # nº de páginas del bono PDF; la landing EN (F3) debe decir lo mismo
# T6: LISTA BLANCA. ASCII, letras Latin-1 (À-ÿ salvo × y ÷, que van en la lista) y exactamente estos símbolos.
# Todo lo demás (CJK, cirílico, puntuación de ancho completo o CJK, ’ curva…) falla.
SIMBOLOS_OK = set('→▸×—÷·−…©£≈✓⅓½⅔¾Σ') | {'\U0001F4F7', '\U0001F4CB'}      # + 📷 📋


def no_permitidos(t):
    return sorted({ch for ch in t if ord(ch) > 127 and ch not in SIMBOLOS_OK
                   and not ('\u00c0' <= ch <= '\u00ff' and ch not in '×÷')})


# T2: IDs internos de la SPEC o del proceso que no deben llegar al comprador.
RX_ID_INTERNO = re.compile(r'\((?:D|E|M|T)\d{1,2}\)|\bR2T?-\d+|\[(?:estimado|derivado)\]|\bSPEC\b|\bv1\.\d\b')
FUNCIONES_SIN_PYCEL = ('IRR', 'XIRR', 'PMT', 'COUNTA')
RX_SIN_PYCEL = re.compile(r'\b(' + '|'.join(FUNCIONES_SIN_PYCEL) + r')\s*\(', re.I)
ERRORES = ('#REF!', '#VALUE!', '#NAME?', '#DIV/0!', '#N/A', '#NUM!', '#NULL!')


# ==========================================================================
# Utilidades
# ==========================================================================
class Resultado:
    def __init__(self):
        self.gates = OrderedDict()
        self.actual = None

    def gate(self, nombre):
        self.actual = nombre
        self.gates.setdefault(nombre, {'fallos': [], 'datos': OrderedDict()})
        return self

    def fallo(self, msg):
        self.gates[self.actual]['fallos'].append(msg)

    def dato(self, k, v):
        self.gates[self.actual]['datos'][k] = v


def cargar(path, data_only=False):
    return openpyxl.load_workbook(path, data_only=data_only)


RX_STR = re.compile(r'"(?:[^"]|"")*"')
RX_REF = re.compile(r"(?:'((?:[^']|'')+)'|([^\W\d][\w.]*))!")
RX_CONV = re.compile(r'<<Conversiones>>\$A\$5:\$B\$\d+')


def normalizar(f, a_es, literales_inv=None):
    """Forma canónica de una fórmula: hojas → «<<nombre ES>>», literales → ES y el final del
    rango de Conversions → «N». `a_es` traduce el nombre de hoja del lado leído al ES."""
    out, pos = [], 0

    def refs(s):
        def rep(m):
            nombre = (m.group(1) if m.group(1) is not None else m.group(2)).replace("''", "'")
            return '<<%s>>' % a_es.get(nombre, '??' + nombre)
        s = RX_REF.sub(rep, s)
        return RX_CONV.sub('<<Conversiones>>$A$5:$B$N', s)
    for m in RX_STR.finditer(f):
        out.append(refs(f[pos:m.start()]))
        lit = m.group(0)[1:-1].replace('""', '"')
        if literales_inv and lit in literales_inv:
            lit = literales_inv[lit]
        out.append('"' + lit + '"')
        pos = m.end()
    out.append(refs(f[pos:]))
    return ''.join(out)


def textos_libro(wb):
    """(hoja, coord, texto) de toda celda de texto que no es fórmula."""
    for ws in wb.worksheets:
        for c in ws._cells.values():
            if isinstance(c.value, str) and c.data_type != 'f':
                yield ws.title, c.coordinate, c.value


def dv_items(wb):
    items = set()
    for ws in wb.worksheets:
        for dv in ws.data_validations.dataValidation:
            f1 = dv.formula1 or ''
            if f1.startswith('"'):
                items.update(f1.strip('"').split(','))
    return items


def xml_parte(path, parte):
    with zipfile.ZipFile(path) as z:
        return z.read(parte).decode('utf-8')


def libros(carpeta):
    """(ES fname, EN fname, ruta ES, ruta EN)."""
    for es in mapas.LIBROS_ES:
        en = mapas.FICHEROS[es]
        yield es, en, os.path.join(ORIGEN, es), os.path.join(carpeta, en)


# ==========================================================================
# Restos de español (gate 2)
# ==========================================================================
SIGNOS_ES = set('¿¡«»€ºª')
RX_TILDE = re.compile(r'[áéíóúüñÁÉÍÓÚÜÑ]')
LISTA_BLANCA_TILDE = {
    'café', 'cafés', 'purée', 'purées', 'puréed', 'sauté', 'sautés', 'sautéed', 'sautéing',
    'açaí', 'jalapeño', 'jalapeños', 'crème', 'brûlée', 'ibérico', 'flambé', 'soufflé',
    'consommé', 'fumé', 'ximénez', 'entremés',
}
FUNCIONALES_ES = {
    'de', 'del', 'la', 'las', 'el', 'los', 'y', 'que', 'con', 'para', 'por', 'una', 'unos',
    'unas', 'es', 'al', 'se', 'su', 'sus', 'como', 'pero', 'muy', 'cuando', 'donde', 'porque',
    'este', 'esta', 'esto', 'estos', 'desde', 'hasta', 'sobre', 'entre', 'tu', 'tus', 'te', 'mi',
    'mis', 'hoja', 'pestaña', 'celda', 'celdas',
}
OFICIO_ES = {
    'iva', 'escandallo', 'escandallos', 'merma', 'mermas', 'albaran', 'albaranes', 'proveedor',
    'proveedores', 'plato', 'platos', 'cocina', 'ventas', 'compras', 'compra', 'semana', 'semanas',
    'coste', 'costes', 'racion', 'raciones', 'pvp', 'tpv', 'euro', 'euros', 'plantilla',
    'plantillas', 'pulpo', 'lubina', 'corvina', 'carrillera', 'solomillo', 'restaurante',
    'inventario', 'factura', 'desproteger', 'revisar', 'cantidad', 'precio', 'precios',
    'unidad', 'unidades', 'docena', 'manojo', 'lata', 'botella', 'objetivo', 'comensal',
    'comensales', 'camarero', 'camareros', 'desperdicio', 'resumen', 'presupuesto',
    'conversiones', 'instrucciones', 'categoria', 'bebida', 'bebidas', 'postre', 'entrante',
    'mantequilla', 'harina', 'azucar', 'huevos', 'nata', 'aceite', 'pescado', 'carne', 'verdura',
    'rendimiento', 'despiece', 'coccion', 'subproducto', 'desecho', 'util', 'alerta',
    'revisa', 'total', 'importe',
}
# «total» e «importe»: «total» es también inglés → se retira abajo; la lista queda explícita.
OFICIO_ES.discard('total')
PROHIBIDAS_EN = {'entrée', 'entree', 'entrées', 'entrees'}
FRASES_BLANCAS = [
    r'\bà la\b', r'\ba la carte\b', r'\bmise en place\b', r'\bpedro xim[ée]nez\b',
    r'\bpico de gallo\b', r'\bde cassis\b', r'\bflor de sal\b', r'\bfleur de sel\b', r'\bherbes de provence\b',
    r'\bcrème de\b', r'\bpâte\b', r'\bpain de mie\b', r'\bdulce de leche\b',
]
RX_FRASES = re.compile('|'.join(FRASES_BLANCAS), re.I)
RX_URL = re.compile(r'https?://\S+|\b[\w-]+(?:\.[\w-]+)*\.(?:es|com|pro|gov|org|uk|co)\b[/\w.#-]*', re.I)
RX_PALABRA = re.compile(r"[A-Za-zÀ-ÖØ-öø-ÿ]+")


def _sin_tilde(p):
    return (p.replace('á', 'a').replace('é', 'e').replace('í', 'i').replace('ó', 'o')
             .replace('ú', 'u').replace('ü', 'u').replace('ñ', 'n'))


def restos_espanol(t):
    """[(motivo, contexto)] de un texto."""
    out = []
    limpio = RX_FRASES.sub(' ', RX_URL.sub(' ', t))
    for ch in limpio:
        if ch in SIGNOS_ES:
            out.append(('signo «%s»' % ch, t[:80]))
    for m in RX_PALABRA.finditer(limpio):
        w = m.group(0)
        wl = w.lower()
        ctx = limpio[max(0, m.start() - 30):m.end() + 30].replace('\n', ' ')
        if wl in PROHIBIDAS_EN:
            out.append(('prohibida D8 «%s»' % w, ctx))
        elif RX_TILDE.search(w) and wl not in LISTA_BLANCA_TILDE:
            out.append(('tilde/eñe «%s»' % w, ctx))
        elif wl in FUNCIONALES_ES:
            out.append(('funcional ES «%s»' % w, ctx))
        elif _sin_tilde(wl) in OFICIO_ES and wl not in LISTA_BLANCA_TILDE:
            out.append(('oficio ES «%s»' % w, ctx))
    return out


# ==========================================================================
# Gate 1 — paridad estructural
# ==========================================================================
def _cf_firma(cf_list, trans):
    """Firma de las reglas CF; `trans` convierte cada fórmula (ES→EN con el Transformador, o identidad)."""
    out = []
    for cf in cf_list:
        for r in cf.rules:
            dxf = r.dxf
            color = None
            if dxf is not None:
                color = (str(dxf.fill.fgColor.rgb) if dxf.fill is not None and dxf.fill.fgColor is not None else None,
                         str(dxf.fill.bgColor.rgb) if dxf.fill is not None and dxf.fill.bgColor is not None else None,
                         str(dxf.font.color.rgb) if dxf.font is not None and dxf.font.color is not None else None)
            cs = None
            if r.colorScale is not None:
                cs = [str(c.rgb) for c in r.colorScale.color]
            out.append((str(cf.sqref), r.type, r.operator, r.priority, bool(r.stopIfTrue),
                        tuple(trans(x) for x in (r.formula or [])), color, cs))
    return sorted(out, key=repr)


def _cf_literales(R, ws, lugar, tot):
    """T1: en cada regla CF con un literal, ese literal tiene que poder salir de las fórmulas de su rango
    (cellIs) o del libro de la hoja (expresión: fórmulas y listas de DV). Si no, el resaltado no salta nunca."""
    lits_hoja = set()
    for c in ws._cells.values():
        if c.data_type == 'f':
            lits_hoja.update(x[1:-1].replace('""', '"') for x in RX_STR.findall(c.value))
    for dv in ws.data_validations.dataValidation:
        if (dv.formula1 or '').startswith('"'):
            lits_hoja.update(dv.formula1.strip('"').split(','))
    for cf in ws.conditional_formatting:
        rango = set()
        for rg in cf.sqref.ranges:
            for fila in ws.iter_rows(min_row=rg.min_row, max_row=rg.max_row, min_col=rg.min_col, max_col=rg.max_col):
                for c in fila:
                    if c.data_type == 'f':
                        rango.update(x[1:-1].replace('""', '"') for x in RX_STR.findall(c.value))
        for regla in cf.rules:
            for f in regla.formula or []:
                for lit in (x[1:-1].replace('""', '"') for x in RX_STR.findall(f)):
                    if lit == '':
                        continue
                    tot['cf_literales'] += 1
                    fuente = rango if regla.type == 'cellIs' else lits_hoja
                    if lit not in fuente:
                        R.fallo('%s: la CF %s busca %r y ninguna fórmula de su rango lo devuelve'
                                % (lugar, cf.sqref, lit))


def _dv_firma(dv):
    return (str(dv.sqref), dv.type, dv.operator, bool(dv.allow_blank), bool(dv.showDropDown),
            bool(dv.showErrorMessage), bool(dv.showInputMessage), dv.errorStyle)


def _prot(ws):
    p = ws.protection
    return tuple(getattr(p, k) for k in ('sheet', 'password', 'sort', 'autoFilter', 'formatColumns',
                                         'formatRows', 'insertRows', 'deleteRows',
                                         'selectLockedCells', 'selectUnlockedCells', 'objects',
                                         'scenarios', 'formatCells', 'insertColumns',
                                         'deleteColumns', 'insertHyperlinks', 'pivotTables'))


def _pagina(ws):
    ps = ws.page_setup
    fit = ws.sheet_properties.pageSetUpPr.fitToPage if ws.sheet_properties.pageSetUpPr else None
    return (ps.orientation, ps.fitToWidth, ps.fitToHeight, fit, str(ws.print_title_rows),
            str(ws.print_title_cols))


def _area(ws):
    pa = ws.print_area
    return pa.split('!')[-1] if pa else None


def gate1(R, carpeta, datos):
    R.gate('1 paridad estructural')
    fin = datos['fin_conv']
    tot = Counter()
    for es, en, pes, pen in libros(carpeta):
        if not os.path.isfile(pen):
            R.fallo('%s: no existe' % en)
            continue
        wes, wen = cargar(pes), cargar(pen)
        esperado = [mapas.hoja_en(h) for h in wes.sheetnames]
        if wen.sheetnames != esperado:
            R.fallo('%s: hojas %r ≠ %r' % (en, wen.sheetnames, esperado))
            continue
        a_es_es = {h: h for h in wes.sheetnames}
        a_es_en = {mapas.hoja_en(h): h for h in wes.sheetnames}
        # T1: sentido DIRECTO y exacto. La fórmula ES pasada por el mismo Transformador que usa aplicar_en tiene
        # que ser IGUAL a la EN: un literal que siga en español («revisa unidades», «ALERTA») ya no se esconde.
        T = aplicar_en.Transformador(OrderedDict((h, mapas.hoja_en(h)) for h in wes.sheetnames), fin)

        def directa(f):
            try:
                return '=' + T.formula(f[1:]) if f.startswith('=') else T.formula(f)
            except aplicar_en.Aborta as e:
                return '<<%s>>' % e
        for h_es in wes.sheetnames:
            ws_es, ws_en = wes[h_es], wen[mapas.hoja_en(h_es)]
            lugar = '%s:%s' % (en, ws_en.title)
            f_es = {c.coordinate: c.value for c in ws_es._cells.values() if c.data_type == 'f'}
            f_en = {c.coordinate: c.value for c in ws_en._cells.values() if c.data_type == 'f'}
            extra_conv = set()
            if h_es == 'Conversiones':                  # E1: fórmulas de bottle/can generadas
                extra_conv = {'B%d' % r['fila'] for r in datos['conv'] if isinstance(r['valor'], str)}
                for k in extra_conv:
                    if f_en.get(k) != next(r['valor'] for r in datos['conv'] if 'B%d' % r['fila'] == k):
                        R.fallo('%s!%s: fórmula E1 ≠ mapas.conversions()' % (lugar, k))
            if set(f_es) != set(f_en) - extra_conv:
                R.fallo('%s: celdas con fórmula distintas: solo ES %s · solo EN %s'
                        % (lugar, sorted(set(f_es) - set(f_en))[:5], sorted(set(f_en) - set(f_es) - extra_conv)[:5]))
            for k, fe in f_es.items():
                fn = f_en.get(k)
                if fn is None:
                    continue
                tot['formulas'] += 1
                if h_es == 'Formatos de Compra' and k.startswith('D') and 5 <= int(k[1:]) <= 16:
                    r = int(k[1:])                       # E3: ml y ×1000
                    if fe != '=IFERROR(ROUND(C%d*100/B%d,4),"")' % (r, r) or \
                            fn != '=IFERROR(ROUND(C%d*1000/B%d,4),"")' % (r, r):
                        R.fallo('%s!%s: E3 %r → %r' % (lugar, k, fe, fn))
                    tot['E3'] += 1
                    continue
                esperado = directa(fe)
                if esperado != fn:
                    R.fallo('%s!%s: fórmula EN %r ≠ ES transformada %r' % (lugar, k, fn[:80], esperado[:80]))
                else:
                    tot['formulas_ok'] += 1
                if 'Conversiones!$A$5:$B$' in fe:
                    tot['vlookup_conv'] += 1
                    if 'Conversions!$A$5:$B$%d,' % fin not in fn:
                        R.fallo('%s!%s: rango de Conversions ≠ $B$%d (E2)' % (lugar, k, fin))
            # DV
            dves = sorted(ws_es.data_validations.dataValidation, key=lambda d: str(d.sqref))
            dven = sorted(ws_en.data_validations.dataValidation, key=lambda d: str(d.sqref))
            if len(dves) != len(dven):
                R.fallo('%s: %d DV en ES y %d en EN' % (lugar, len(dves), len(dven)))
            for de, dn in zip(dves, dven):
                tot['dv'] += 1
                if _dv_firma(de) != _dv_firma(dn):
                    R.fallo('%s: DV %s cambia (%r → %r)' % (lugar, de.sqref, _dv_firma(de), _dv_firma(dn)))
                f1e, f1n = de.formula1 or '', dn.formula1 or ''
                if f1e == aplicar_en.UNIDADES_ES_DV:
                    tot['E5'] += 1
                    if f1n != mapas.DV_UNIDADES_EN:
                        R.fallo('%s: DV %s no es la lista D7 (E5)' % (lugar, dn.sqref))
                elif f1e == aplicar_en.DV12_ES:
                    if f1n != mapas.DV12_EN:
                        R.fallo('%s: DV %s ≠ tipos EN del 12' % (lugar, dn.sqref))
                elif f1e == aplicar_en.MOTIVOS_ES:
                    if f1n != mapas.DV_MOTIVOS_EN:
                        R.fallo('%s: DV %s ≠ motivos EN' % (lugar, dn.sqref))
                elif f1e.startswith('"'):
                    if f1n != f1e:
                        R.fallo('%s: DV %s lista literal cambiada' % (lugar, dn.sqref))
                elif T.refs(f1e) != f1n:
                    R.fallo('%s: DV %s %r ≠ %r' % (lugar, dn.sqref, T.refs(f1e), f1n))
                if (de.formula2 or dn.formula2) and T.refs(de.formula2 or '') != (dn.formula2 or ''):
                    R.fallo('%s: DV %s formula2 %r ≠ %r' % (lugar, dn.sqref, de.formula2, dn.formula2))
            # CF
            cfe = _cf_firma(ws_es.conditional_formatting, T.formula)
            cfn = _cf_firma(ws_en.conditional_formatting, lambda x: x)
            tot['cf'] += len(cfe)
            if cfe != cfn:
                dif = [(a, b) for a, b in zip(cfe, cfn) if a != b] or [(cfe[:1], cfn[:1])]
                R.fallo('%s: formato condicional distinto (ES transformado → EN): %r' % (lugar, dif[0]))
            _cf_literales(R, ws_en, lugar, tot)
            # merges, protección, paneles, impresión, ajuste de página, autofiltro
            if sorted(map(str, ws_es.merged_cells.ranges)) != sorted(map(str, ws_en.merged_cells.ranges)):
                R.fallo('%s: combinaciones distintas' % lugar)
            tot['merges'] += len(ws_es.merged_cells.ranges)
            if _prot(ws_es) != _prot(ws_en):
                R.fallo('%s: protección distinta' % lugar)
            if ws_es.freeze_panes != ws_en.freeze_panes:
                R.fallo('%s: paneles %s ≠ %s' % (lugar, ws_es.freeze_panes, ws_en.freeze_panes))
            if _pagina(ws_es) != _pagina(ws_en):
                R.fallo('%s: ajuste o títulos de impresión %r ≠ %r' % (lugar, _pagina(ws_es), _pagina(ws_en)))
            if str(ws_es.auto_filter.ref) != str(ws_en.auto_filter.ref):
                R.fallo('%s: autofiltro distinto' % lugar)
            ae, an = _area(ws_es), _area(ws_en)
            if h_es == 'Conversiones':
                if an != '$A$1:$J$%d' % fin:
                    R.fallo('%s: área de impresión %s ≠ $A$1:$J$%d (E2)' % (lugar, an, fin))
            elif ae != an:
                R.fallo('%s: área de impresión %s ≠ %s' % (lugar, ae, an))
            # celdas desbloqueadas (entradas) — Conversions no está protegida (E1)
            if h_es != 'Conversiones':
                ue = {c.coordinate for c in ws_es._cells.values() if c.protection.locked is False}
                un = {c.coordinate for c in ws_en._cells.values() if c.protection.locked is False}
                if ue != un:
                    R.fallo('%s: celdas desbloqueadas distintas (solo ES %s, solo EN %s)'
                            % (lugar, sorted(ue - un)[:5], sorted(un - ue)[:5]))
                tot['desbloqueadas'] += len(ue)
        # gráficos
        ge, gn = extraer_textos.graficos(pes), extraer_textos.graficos(pen)
        if len(ge) != len(gn):
            R.fallo('%s: %d gráficos ES y %d EN' % (en, len(ge), len(gn)))
        for a, b in zip(ge, gn):
            tot['graficos'] += 1
            if mapas.hoja_en(a['hoja']) != b['hoja'] or a['tipos'] != b['tipos']:
                R.fallo('%s: gráfico %s ancla/tipo distintos' % (en, b['parte']))
            if [normalizar(x, a_es_es) for x in a['refs']] != [normalizar(x, a_es_en) for x in b['refs']]:
                R.fallo('%s: series del gráfico distintas: %r' % (en, b['refs']))
            if '<style val="2"/>' in xml_parte(pes, a['parte']) and '<style val="2"/>' not in xml_parte(pen, b['parte']):
                R.fallo('%s: el gráfico perdió su estilo' % en)
            for h in b['hojas_citadas']:
                if h not in wen.sheetnames:
                    R.fallo('%s: el gráfico cita la hoja inexistente %r' % (en, h))
        dib_es = sorted(n for n in zipfile.ZipFile(pes).namelist() if n.startswith('xl/drawings/drawing'))
        dib_en = sorted(n for n in zipfile.ZipFile(pen).namelist() if n.startswith('xl/drawings/drawing'))
        if [xml_parte(pes, n) for n in dib_es] != [xml_parte(pen, n) for n in dib_en]:
            R.fallo('%s: el ancla o el tamaño de los gráficos (drawings) cambia' % en)
        # ninguna referencia a una hoja inexistente (fórmulas, DV, áreas)
        for ws in wen.worksheets:
            for c in ws._cells.values():
                if c.data_type == 'f':
                    for h in extraer_textos.refs_hoja(c.value):
                        if h not in wen.sheetnames:
                            R.fallo('%s:%s!%s cita la hoja inexistente %r' % (en, ws.title, c.coordinate, h))
            for dv in ws.data_validations.dataValidation:
                for h in extraer_textos.refs_hoja(dv.formula1 or ''):
                    if h not in wen.sheetnames:
                        R.fallo('%s:%s DV cita la hoja inexistente %r' % (en, ws.title, h))
    for k, v in sorted(tot.items()):
        R.dato(k, v)


# ==========================================================================
# Gate 2 — restos y caracteres
# ==========================================================================
def piezas_de_texto(pen):
    """Todo lo que lleva texto en un libro EN: (tipo, lugar, texto)."""
    wb = cargar(pen)
    for h, c, v in textos_libro(wb):
        yield 'valor', '%s!%s' % (h, c), v
    for ws in wb.worksheets:
        yield 'hoja', ws.title, ws.title
        for c in ws._cells.values():
            if isinstance(c.number_format, str) and c.number_format != 'General':
                yield 'formato', '%s!%s' % (ws.title, c.coordinate), c.number_format
        for dv in ws.data_validations.dataValidation:
            for attr in ('error', 'errorTitle', 'prompt', 'promptTitle', 'formula1'):
                v = getattr(dv, attr)
                if v and (attr != 'formula1' or v.startswith('"')):
                    yield 'dv', '%s DV %s %s' % (ws.title, dv.sqref, attr), v
        for parte in ('oddHeader', 'oddFooter', 'evenHeader', 'evenFooter', 'firstHeader', 'firstFooter'):
            hf = getattr(ws, parte)
            for pos in ('left', 'center', 'right'):
                t = getattr(hf, pos).text
                if t:
                    yield 'cabecera-pie', '%s %s.%s' % (ws.title, parte, pos), t
    p = wb.properties
    for k in ('title', 'subject', 'creator', 'keywords', 'description', 'category', 'lastModifiedBy'):
        v = getattr(p, k)
        if v:
            yield 'docprops', k, v
    for ws in wb.worksheets:                     # T1: literales de fórmula y de CF
        for c in ws._cells.values():
            if c.data_type == 'f':
                for x in RX_STR.findall(c.value):
                    yield 'literal', '%s!%s' % (ws.title, c.coordinate), x[1:-1].replace('""', '"')
        for cf in ws.conditional_formatting:
            for regla in cf.rules:
                for f in regla.formula or []:
                    for x in RX_STR.findall(f):
                        yield 'literal', '%s CF %s' % (ws.title, cf.sqref), x[1:-1].replace('""', '"')
    for g in extraer_textos.graficos(pen):
        if g['titulo']:
            yield 'grafico', g['parte'], g['titulo']
        for t in g['textos_cache']:
            yield 'grafico', g['parte'] + ' caché', t
    estilos = xml_parte(pen, 'xl/styles.xml')
    for m in re.finditer(r'<numFmt [^>]*formatCode="([^"]*)"', estilos):
        yield 'styles.xml', 'numFmt', html.unescape(m.group(1))


def gate2(R, carpeta, datos):
    R.gate('2 restos y caracteres')
    n = Counter()
    for es, en, pes, pen in libros(carpeta):
        for tipo, lugar, t in piezas_de_texto(pen):
            n[tipo] += 1
            donde = '%s %s' % (en, lugar)
            if '€' in t:
                R.fallo('%s: «€» en %s: %r' % (donde, tipo, t[:60]))
            for sim in ('$', '£'):                     # D5: solo como importe dentro de un texto
                if sim not in t:
                    continue
                if tipo != 'valor':
                    R.fallo('%s: «%s» en %s: %r' % (donde, sim, tipo, t[:60]))
                elif len(re.findall(re.escape(sim) + r'\d', t)) != t.count(sim):
                    R.fallo('%s: «%s» que no es un importe: %r' % (donde, sim, t[:80]))
                else:
                    n['moneda_en_texto_' + sim] += 1
            malos = no_permitidos(t)
            if malos:
                R.fallo('%s: carácter fuera de la lista blanca %r en %s: %r' % (donde, ''.join(malos), tipo, t[:60]))
            if tipo not in ('formato', 'styles.xml'):
                m = RX_ID_INTERNO.search(t)
                if m:
                    R.fallo('%s: ID interno %r en %s: …%s…' % (donde, m.group(0), tipo,
                                                               t[max(0, m.start() - 40):m.end() + 10]))
            if tipo == 'formato' or tipo == 'styles.xml':
                if re.search(r'dd/mm', t, re.I):
                    R.fallo('%s: formato dd/mm %r' % (donde, t))
                continue
            for motivo, ctx in restos_espanol(t):
                R.fallo('%s: %s · …%s…' % (donde, motivo, ctx))
    for k, v in sorted(n.items()):
        R.dato(k, v)


# ==========================================================================
# Gate 3 — cálculo
# ==========================================================================
def _num(v):
    return isinstance(v, (int, float)) and not isinstance(v, bool)


def celdas_sin_v(path):
    """T5: {(hoja, celda)} de toda celda con <f> y SIN <v> en el XML. Una vacía legítima lleva <v></v>;
    openpyxl lee las dos como None, así que se mira el XML."""
    fuera = set()
    with zipfile.ZipFile(path) as z:
        wbx = z.read('xl/workbook.xml').decode('utf-8')
        rels = z.read('xl/_rels/workbook.xml.rels').decode('utf-8')
        destino = {}
        for tag in re.findall(r'<Relationship\b[^>]*>', rels):
            i, t = re.search(r'\bId="([^"]+)"', tag), re.search(r'\bTarget="([^"]+)"', tag)
            if i and t:
                destino[i.group(1)] = t.group(1)
        for tag in re.findall(r'<sheet\b[^>]*>', wbx):
            nombre = html.unescape(re.search(r'\bname="([^"]*)"', tag).group(1))
            rid = re.search(r'\br:id="([^"]+)"', tag).group(1)
            t = destino[rid]
            parte = t.lstrip('/') if t.startswith('/') else 'xl/' + t
            xml = z.read(parte).decode('utf-8')
            for m in re.finditer(r'<c\b([^>]*?)(?:/>|>(.*?)</c>)', xml, re.S):
                dentro = m.group(2) or ''
                if '<f' in dentro and '<v' not in dentro:
                    fuera.add((nombre, re.search(r'\br="([A-Z]+\d+)"', m.group(1)).group(1)))
    return fuera


def gate3(R, carpeta, datos):
    R.gate('3 cálculo')
    n = Counter()
    for es, en, pes, pen in libros(carpeta):
        wf, wv = cargar(pen), cargar(pen, data_only=True)
        sin_v = celdas_sin_v(pen)
        n['formulas_sin_v_en_xml'] += len(sin_v)
        for ws in wf.worksheets:
            wsv = wv[ws.title]
            for c in ws._cells.values():
                if c.data_type != 'f':
                    continue
                n['formulas'] += 1
                v = wsv[c.coordinate].value
                if RX_SIN_PYCEL.search(c.value):
                    R.fallo('%s:%s!%s usa una función que pycel no evalúa' % (en, ws.title, c.coordinate))
                if (ws.title, c.coordinate) in sin_v:
                    R.fallo('%s:%s!%s sin caché (<f> sin <v> en el XML)' % (en, ws.title, c.coordinate))
                    continue
                if v is None:
                    n['vacias_legitimas'] += 1            # <v></v>: resultado "" guardado
                    continue
                if isinstance(v, str) and v in ERRORES:
                    if v == '#N/A' and 'NA()' in c.value:
                        n['na_de_grafico'] += 1
                    else:
                        R.fallo('%s:%s!%s = %s' % (en, ws.title, c.coordinate, v))
    for k, v in sorted(n.items()):
        R.dato(k, v)
    checks_conversions(R, carpeta, datos)
    sensibilidad(R, carpeta, datos)


def leer_conversions(ws, wsv):
    filas = OrderedDict()
    r = mapas.FILA0
    while ws.cell(r, 1).value not in (None, ''):
        filas[ws.cell(r, 1).value] = {'fila': r, 'valor': wsv.cell(r, 2).value, 'formula': ws.cell(r, 2).value}
        r += 1
    return filas


def checks_conversions(R, carpeta, datos):
    """Los 6 checks de §2.3 sobre las 8 tablas Conversions EN."""
    lista = set(mapas.UNIDADES_EN)
    fin = datos['fin_conv']
    cat_en = list(mapas.CATEGORIAS.values())
    for es, en, pes, pen in libros(carpeta):
        wf = cargar(pen)
        if 'Conversions' not in wf.sheetnames:
            continue
        wv = cargar(pen, data_only=True)
        conv = leer_conversions(wf['Conversions'], wv['Conversions'])
        claves = list(conv)
        lugar = en + ':Conversions'
        # 1 duplicados (VLOOKUP no distingue mayúsculas)
        if len({k.lower() for k in claves}) != len(claves):
            R.fallo('%s: claves duplicadas' % lugar)
        # 2 unidades en la lista D7
        for k in claves:
            o, _, d = k.partition('→')
            if o not in lista or d not in lista:
                R.fallo('%s: clave %r con unidad fuera de la lista D7' % (lugar, k))
        ks = set(claves)
        # 3 resolución: base↔base, packs, envases; ningún pack como destino salvo identidad
        for dim, tabla in mapas.BASE.items():
            for o in tabla:
                for d in tabla:
                    if o + '→' + d not in ks:
                        R.fallo('%s: falta base↔base %s→%s' % (lugar, o, d))
        for nombre, dim, tam, subs, bebida, *_ in mapas.PACKS:
            for d in [nombre] + list(subs) + [u for u in mapas.BASE[dim]
                                                if not (bebida and u in mapas.EXCLUIDOS_BEBIDA)]:
                if nombre + '→' + d not in ks:
                    R.fallo('%s: falta %s→%s' % (lugar, nombre, d))
            if bebida:
                for u in mapas.EXCLUIDOS_BEBIDA:
                    if nombre + '→' + u in ks:
                        R.fallo('%s: %s→%s no debe existir (D7)' % (lugar, nombre, u))
        for env in mapas.ENVASES:
            for d in mapas.DESTINOS_ENVASE + [env]:
                if env + '→' + d not in ks:
                    R.fallo('%s: falta %s→%s' % (lugar, env, d))
        for k in claves:
            o, _, d = k.partition('→')
            if d in mapas.PACKS_NOMBRES and o != d:
                R.fallo('%s: pack como destino en %s' % (lugar, k))
            dims = set()
            for u in (o, d):
                try:
                    dims.add(mapas.dimension(u))
                except KeyError:
                    pass
            if dims == {'masa', 'volumen'}:
                R.fallo('%s: peso↔volumen en %s' % (lugar, k))
        # valores = factores de mapas (en caché)
        factores = {r['clave']: r['factor'] for r in datos['conv']}
        for k, info in conv.items():
            v = info['valor']
            if not _num(v) or abs(v - factores.get(k, float('nan'))) > 1e-6 * max(1, abs(v)):
                R.fallo('%s: %s = %r ≠ %r' % (lugar, k, v, factores.get(k)))
        # 4 envases: todas sus claves de volumen salen de su única celda de tamaño
        for env in mapas.ENVASES:
            tam = conv.get(env + '→ml')
            if not tam or isinstance(tam['formula'], str):
                R.fallo('%s: %s→ml no es una celda de valor' % (lugar, env))
                continue
            celda = 'B%d' % tam['fila']
            for d in ('fl oz', 'cl', 'L'):
                f = conv.get(env + '→' + d, {}).get('formula')
                if not (isinstance(f, str) and re.search(r'\b%s\b' % celda, f)):
                    R.fallo('%s: %s→%s no sale de %s' % (lugar, env, d, celda))
        # 5 filas de ejemplo: factor y categoría resueltos
        for ws in wf.worksheets:
            if ws.title in ('Conversions', 'Trim Loss Factors', 'Instructions'):
                continue
            wsv = wv[ws.title]
            dv = [d for d in ws.data_validations.dataValidation if d.formula1 == mapas.DV_UNIDADES_EN]
            if not dv:
                continue
            rng = [x for x in dv[0].sqref.ranges if x.min_col == 3][0]
            for r in range(rng.min_row, rng.max_row + 1):
                if ws.cell(r, 1).value in (None, ''):
                    continue
                g = wsv.cell(r, 7).value
                if not _num(g):
                    R.fallo('%s:%s!G%d: factor %r sin resolver (%s→%s)'
                            % (en, ws.title, r, g, ws.cell(r, 3).value, ws.cell(r, 6).value))
                if ws.cell(r, 2).value not in cat_en:
                    R.fallo('%s:%s!B%d: categoría %r' % (en, ws.title, r, ws.cell(r, 2).value))
                if not _num(wsv.cell(r, 8).value):
                    R.fallo('%s:%s!H%d: merma sin valor' % (en, ws.title, r))
                if not _num(wsv.cell(r, 10).value):
                    R.fallo('%s:%s!J%d: coste %r' % (en, ws.title, r, wsv.cell(r, 10).value))
                R.gates[R.actual]['datos']['filas_ejemplo'] = R.gates[R.actual]['datos'].get('filas_ejemplo', 0) + 1
    R.dato('conversions_n', datos['n'])


def _pycel(path):
    from pycel import ExcelCompiler
    return ExcelCompiler(filename=path)


def _ev(xl, ref):
    with open(os.devnull, 'w') as dn, contextlib.redirect_stderr(dn):
        v = xl.evaluate(ref)
    try:
        import numpy as np
        if isinstance(v, np.generic):
            v = v.item()
    except ImportError:
        pass
    return v


def sensibilidad(R, carpeta, datos):
    """Check 6 de §2.3 (una pareja nueva en la primera fila libre se resuelve) y sensibilidad:
    un precio, una cantidad y el impuesto mueven lo que deben (pycel sobre el 01 EN)."""
    pen = os.path.join(carpeta, mapas.FICHEROS['01-escandallo-estandar.xlsx'])
    xl = _pycel(pen)
    H = "'Recipe Cost Card'!"
    j5, j30, j33 = _ev(xl, H + 'J5'), _ev(xl, H + 'J30'), _ev(xl, H + 'J33')
    _ev(xl, H + 'J35')                                  # mete I34 en el mapa de pycel
    g15 = _ev(xl, H + 'G15')                            # y C15/F15 + el rango de Conversions
    d5 = _ev(xl, H + 'D5')
    xl.set_value(H + 'D5', d5 * 1.10)
    j5b, j30b = _ev(xl, H + 'J5'), _ev(xl, H + 'J30')
    if abs(j5b - j5 * 1.10) > 1e-6 or not abs((j30b - j30) - (j5b - j5) * 1.10) < 1e-6:
        R.fallo('sensibilidad: +10 %% en D5 da J5 %.4f→%.4f y J30 %.4f→%.4f' % (j5, j5b, j30, j30b))
    xl.set_value(H + 'D5', d5)
    xl.set_value(H + 'I34', 0.2)
    j35 = _ev(xl, H + 'J35')
    if abs(j35 - j33 * 1.2) > 1e-6:
        R.fallo('sensibilidad: impuesto 20 %% da J35 %.4f (≠ %.4f)' % (j35, j33 * 1.2))
    xl.set_value(H + 'I34', 0)
    e5 = _ev(xl, H + 'E5')
    xl.set_value(H + 'E5', e5 * 2)
    if abs(_ev(xl, H + 'J5') - j5 * 2) > 1e-6:
        R.fallo('sensibilidad: doble cantidad no dobla el coste de la fila')
    xl.set_value(H + 'E5', e5)
    # T7: celda de tamaño del envase vacía o a 0 → Factor «?» y Cost «check units» (nunca un coste en blanco)
    tam = next(r['fila'] for r in datos['conv'] if r['clave'] == 'bottle→ml')
    _ev(xl, H + 'G16')                                  # mete la fila 16 y el rango de Conversions en pycel
    _ev(xl, H + 'J16')
    xl.set_value(H + 'H16', 0)
    xl.set_value(H + 'C16', 'bottle')
    xl.set_value(H + 'F16', 'fl oz')
    xl.set_value(H + 'E16', 1)
    xl.set_value(H + 'D16', 10)
    g_ok = _ev(xl, H + 'G16')
    if not (_num(g_ok) and abs(g_ok - mapas.ENVASES['bottle']['ml'] / float(mapas.FLOZ_ML)) < 1e-6):
        R.fallo('T7: bottle→fl oz con la celda de tamaño llena da G16 %r' % g_ok)
    # pycel no recalcula con set_value(None): la celda borrada se prueba como 0 (lo que Excel lee de una vacía
    # en una división), como texto vacío y como texto.
    for vacio in (0, '', 'n/a'):
        xl.set_value('Conversions!B%d' % tam, vacio)
        g16, j16 = _ev(xl, H + 'G16'), _ev(xl, H + 'J16')
        if (g16, j16) != ('?', 'check units'):
            R.fallo('T7: celda de tamaño vacía (B%d = %r): G16 %r y J16 %r (se esperaba "?" y "check units")'
                    % (tam, vacio, g16, j16))
    xl.set_value('Conversions!B%d' % tam, mapas.ENVASES['bottle']['ml'])
    for col in 'CDEF':
        xl.set_value(H + '%s16' % col, None)
    # check 6: pareja nueva (bunch→g) en la primera fila libre de Conversions
    libre = mapas.FILA0 + datos['n']
    xl.set_value(H + 'C15', 'bunch')
    xl.set_value(H + 'F15', 'g')
    xl.set_value(H + 'E15', 10)
    antes = _ev(xl, H + 'G15')
    xl.set_value("Conversions!A%d" % libre, 'bunch→g')
    xl.set_value("Conversions!B%d" % libre, 30)
    despues = _ev(xl, H + 'G15')
    if antes != '?' or despues != 30:
        R.fallo('check 6: pareja nueva en Conversions!A%d: G15 %r → %r (se esperaba "?" → 30)'
                % (libre, antes, despues))
    ultima = datos['fin_conv']
    xl.set_value("Conversions!A%d" % libre, None)
    xl.set_value("Conversions!B%d" % libre, None)
    xl.set_value("Conversions!A%d" % ultima, 'bunch→g')
    xl.set_value("Conversions!B%d" % ultima, 25)
    if _ev(xl, H + 'G15') != 25:
        R.fallo('check 6: la última fila libre (A%d) no entra en el VLOOKUP' % ultima)
    R.dato('sensibilidad', 'precio, cantidad, impuesto, celda de tamaño vacía (T7) y fila libre %d-%d'
           % (libre, ultima))


# ==========================================================================
# Gate 4 — ejemplos creíbles
# ==========================================================================
def gate4(R, carpeta, datos):
    R.gate('4 ejemplos creíbles')
    r = subprocess.run([sys.executable, VALIDAR_MERCADO], capture_output=True, text=True, cwd=AQUI)
    ult = [l for l in r.stdout.strip().splitlines() if l.strip()][-3:]
    R.dato('validar_mercado', ' | '.join(ult))
    if r.returncode != 0:
        R.fallo('validar_mercado.py sale con %d: %s' % (r.returncode, ' | '.join(ult)))
    M = datos['mercado']
    d12 = M['d12']
    menus = defaultdict(float)
    for rec in M['recetas']:
        pen = os.path.join(carpeta, mapas.FICHEROS[[f for f in mapas.LIBROS_ES if f.startswith(rec['libro'])][0]])
        wv = datos['cache_v'].setdefault(pen, cargar(pen, data_only=True))
        wf = datos['cache_f'].setdefault(pen, cargar(pen))
        ws, wsv = wf[rec['hoja_en']], wv[rec['hoja_en']]
        fila_cpr = fila_fc = None
        for rr in range(1, ws.max_row + 1):
            a = ws.cell(rr, 1).value
            if isinstance(a, str) and a.upper().startswith('COST PER') and fila_cpr is None:
                fila_cpr = rr
            if isinstance(a, str) and a.lower().startswith('actual food cost') or \
                    (isinstance(a, str) and a.lower().startswith('actual pour cost')):
                fila_fc = rr
        if not fila_cpr or not fila_fc:
            R.fallo('%s: sin fila de coste por ración o de food cost real' % rec['hoja_en'])
            continue
        cpr = wsv.cell(fila_cpr, 10).value
        fc = wsv.cell(fila_fc, 10).value
        precio = ws.cell(fila_fc - 1, 9).value
        if rec['libro'] in ('02', '03'):
            # CHEF-03: el pase no se vende suelto → «Current menu price» vacío; el food cost es el del menú
            if precio not in (None, ''):
                R.fallo('%s: el pase lleva precio propio (%r); debe ir vacío' % (rec['hoja_en'], precio))
            if not _num(cpr) or abs(cpr - rec['coste_racion']) > TOL * rec['coste_racion']:
                R.fallo('%s: coste por ración del xlsx %r ≠ mercado %.4f' % (rec['hoja_en'], cpr, rec['coste_racion']))
            menus[rec['libro']] += cpr if _num(cpr) else 0
            continue
        if not (_num(cpr) and _num(fc)):
            R.fallo('%s: coste %r o food cost real %r sin valor' % (rec['hoja_en'], cpr, fc))
            continue
        if abs(cpr - rec['coste_racion']) > TOL * rec['coste_racion']:
            R.fallo('%s: coste por ración del xlsx %.4f ≠ mercado %.4f' % (rec['hoja_en'], cpr, rec['coste_racion']))
        if precio != rec['precio_carta_ref']:
            R.fallo('%s: precio de carta precargado %r ≠ %r' % (rec['hoja_en'], precio, rec['precio_carta_ref']))
        lo, hi = d12[rec['tipo_local']]['min'], d12[rec['tipo_local']]['max']
        if rec['libro'] == '04':
            lo, hi = 0.18, 0.24
        if not (lo - 1e-9 <= fc <= hi + 1e-9):
            R.fallo('%s: food cost implícito %.1f %% fuera de %s (%d-%d %%)'
                    % (rec['hoja_en'], fc * 100, rec['tipo_local'], lo * 100, hi * 100))
        R.dato(rec['libro'] + ' ' + rec['hoja_en'], '%.1f %%' % (fc * 100))
    # menús 02 y 03: a nivel de menú (el precio de cada pase es su parte del menú)
    for libro, clave in (('02', 'menu_02'), ('03', 'menu_03')):
        m = M[clave]
        extras = sum(x['valor'] for x in m.get('extras', {}).values())
        fc = (menus[libro] + extras) / m['precio_carta_ref']
        lo, hi = d12[m['tipo_local']]['min'], d12[m['tipo_local']]['max']
        R.dato('menú ' + libro, '%.1f %%' % (fc * 100))
        if not (lo <= fc <= hi):
            R.fallo('menú %s: food cost %.1f %% fuera de D12 %s' % (libro, fc * 100, m['tipo_local']))
        # la nota del Summary cita el precio del menú y el food cost que sale del xlsx
        pen = os.path.join(carpeta, mapas.FICHEROS[[f for f in mapas.LIBROS_ES if f.startswith(libro)][0]])
        hoja = 'Summary' if libro == '02' else 'Menu Summary'
        nota = datos['cache_f'].setdefault(pen, cargar(pen))[hoja][m['nota_resumen']['celda']].value or ''
        if nota != m['nota_resumen']['texto'] or '{:.1f}%'.format(fc * 100) not in nota \
                or '${:,.0f}'.format(m['precio_carta_ref']) not in nota:
            R.fallo('menú %s: la nota %s!%s no cita el precio %s y el food cost %.1f%% (%r)'
                    % (libro, hoja, m['nota_resumen']['celda'], m['precio_carta_ref'], fc * 100, nota[:90]))


# ==========================================================================
# Gate 5 — formato
# ==========================================================================
def gate5(R, carpeta, datos, mes):
    R.gate('5 formato')
    censo = json.load(open(os.path.join(AQUI, 'censo_es.json'), encoding='utf-8'))
    linea = aplicar_en.linea_version(mes)
    n = Counter()
    for es, en, pes, pen in libros(carpeta):
        wb = cargar(pen)
        for ws in wb.worksheets:
            n['hojas'] += 1
            if ws.page_setup.paperSize != PAPEL_LETTER:
                R.fallo('%s:%s papel %r ≠ Letter' % (en, ws.title, ws.page_setup.paperSize))
            if ws.oddFooter.center.text != PIE_EN:
                R.fallo('%s:%s pie %r' % (en, ws.title, ws.oddFooter.center.text))
            for c in ws._cells.values():
                f = c.number_format or ''
                if '€' in f or '$' in f or re.search(r'dd/mm', f, re.I):
                    R.fallo('%s:%s!%s formato %r' % (en, ws.title, c.coordinate, f))
        # D6 (CHEF-01): en toda fila rotulada «Tax rate (%)», cualquier celda numérica escrita vale 0
        for ws in wb.worksheets:
            for c in list(ws._cells.values()):
                if isinstance(c.value, str) and c.data_type != 'f' and c.value.strip().startswith('Tax rate (%)'):
                    n['filas_tax_rate'] += 1
                    for x in [x for (r, _k), x in list(ws._cells.items()) if r == c.row]:   # sin crear celdas
                        if _num(x.value) and x.value != 0:
                            R.fallo('%s:%s!%s = %r en la fila «%s» (D6: 0 %% por defecto)'
                                    % (en, ws.title, x.coordinate, x.value, c.value))
        # fechas: por regla (todas las del censo ES), numFmtId 14
        det = censo['libros'][es]['hojas_detalle']
        for h_es, info in det.items():
            ws = wb[mapas.hoja_en(h_es)]
            for fch in info['fechas']:
                n['fechas'] += 1
                c = ws[fch['celda']]
                if c._style.numFmtId != 14:
                    R.fallo('%s:%s!%s fecha con numFmtId %s' % (en, ws.title, fch['celda'], c._style.numFmtId))
        # fechas que no están en el censo (formato de fecha en otra celda)
        for ws in wb.worksheets:
            for c in ws._cells.values():
                if c._style.numFmtId == 14:
                    n['celdas_numfmt14'] += 1
        # línea de versión D14
        ins = wb['Instructions']
        vers = [c.value for c in ins._cells.values() if isinstance(c.value, str) and c.value.startswith('Version ')]
        if vers != [linea]:
            R.fallo('%s: línea de versión %r ≠ %r' % (en, vers, linea))
        p = wb.properties
        esperado = {'subject': aplicar_en.SUBJECT, 'description': aplicar_en.DESCRIPTION,
                    'creator': 'AI Chef Pro', 'lastModifiedBy': 'AI Chef Pro',
                    'category': 'AI Chef Pro · Digital products'}
        for k, v in esperado.items():
            if getattr(p, k) != v:
                R.fallo('%s: docProps %s = %r' % (en, k, getattr(p, k)))
        if not p.title or not p.title.endswith('· Recipe Costing Kit Pro'):
            R.fallo('%s: docProps title %r' % (en, p.title))
        if not p.keywords or 'recipe costing kit' not in p.keywords.lower():
            R.fallo('%s: docProps keywords %r' % (en, p.keywords))
    if n['fechas'] != censo['totales']['fechas']:
        R.fallo('fechas contadas %d ≠ censo %d' % (n['fechas'], censo['totales']['fechas']))
    if n['celdas_numfmt14'] != n['fechas']:
        R.fallo('%d celdas con numFmtId 14 y %d fechas en el censo' % (n['celdas_numfmt14'], n['fechas']))
    for k, v in sorted(n.items()):
        R.dato(k, v)
    R.dato('version', linea)


# ==========================================================================
# Gate 6 — coherencia de texto acotada
# ==========================================================================
ETIQUETAS_D12 = OrderedDict([
    ('fine_dining', [r'fine[- ]dining']), ('fast_casual', [r'fast[- ]casual']),
    ('casual', [r'(?<!fast )(?<!fast-)casual(?: dining)?']), ('cafe', [r'caf[ée]s?\b']),
    ('catering', [r'catering']), ('food_truck', [r'food trucks?']),
    ('hotel_fb', [r'hotel(?: f&b)?']), ('pastry_bakery', [r'pastry(?: (?:&|and) bakery)?', r'bakery']),
    ('bar', [r'\bbars?\b(?! (?:&|and) cocktails)', r'bar (?:&|and) cocktails']),
    ('delivery', [r'delivery']),
])
RX_RANGO = re.compile(r'(\d{1,2}(?:\.\d)?)\s*(?:%\s*)?[-–]\s*(\d{1,2}(?:\.\d)?)\s*%')
RX_COMILLAS_D = re.compile(r'"([^"\n]{1,60})"')
RX_COMILLAS_S = re.compile(r"(?<![A-Za-z])'([^'\n]{1,40})'(?![A-Za-z])")
# Cita entre comillas que no es pestaña ni rótulo literal: la forma de la clave de Conversions
# («the key is written "purchase→recipe"»), que es el paréntesis del rótulo A4.
CITAS_BLANCAS = {'purchase→recipe'}
RX_LIBRO_NUM = re.compile(r'(?<![\d.,$])\b(0[1-9]|1[0-3])\b(?![.,]\d)')


def textos_citables(wb):
    """Textos donde se citan pestañas o rótulos: celdas, mensajes de DV y cabeceras/pies (T4)."""
    for h, coord, t in textos_libro(wb):
        yield h, coord, t
    for ws in wb.worksheets:
        for dv in ws.data_validations.dataValidation:
            for attr in ('error', 'prompt'):
                v = getattr(dv, attr)
                if v:
                    yield ws.title, 'DV %s %s' % (dv.sqref, attr), v
        for parte in ('oddHeader', 'oddFooter', 'evenHeader', 'evenFooter', 'firstHeader', 'firstFooter'):
            hf = getattr(ws, parte)
            for pos in ('left', 'center', 'right'):
                t = getattr(hf, pos).text
                if t:
                    yield ws.title, '%s.%s' % (parte, pos), t


def cita_casa(q, etiquetas):
    """T4: coincidencia EXACTA (con mayúsculas) contra un rótulo. Un prefijo solo vale si la cita acaba en «…»
    o si el rótulo sigue con «:» o « (» (p. ej. «Current menu price» → «Current menu price (pre-tax / ex VAT)»),
    y la marca «X:» que abre una nota dentro de una celda (13: «Sample data:»)."""
    limpia = q.strip().rstrip('.,;:')
    for e in etiquetas:
        if e == limpia or e.startswith(limpia + ':') or e.startswith(limpia + ' ('):
            return True
        if (limpia + ':') in e:           # marca literal de una nota («Sample data:»), con mayúsculas
            return True
        if limpia.endswith('…') and e.startswith(limpia[:-1]):
            return True
    return False


def libros_nombrados(t):
    """Libros que un texto nombra: «template 04», «09-food-waste-tracker», «Used in: 02 …»,
    «files 01-04», «BONUS» y «recipe cost card(s)» (= 01-08)."""
    out = set(RX_LIBRO_NUM.findall(t))
    if 'BONUS' in t:
        out.add('BONUS')
    if re.search(r'recipe cost cards?', t, re.I):
        out.update('%02d' % i for i in range(1, 9))
    return out


def gate6(R, carpeta, datos):
    R.gate('6 coherencia de texto')
    M = datos['mercado']
    d12 = {k: (round(v['min'] * 100, 1), round(v['max'] * 100, 1)) for k, v in M['d12'].items()}
    blancos = set()
    for x in M['lista_blanca_rangos'] + datos['extra'].get('rangos_para_lista_blanca', []):
        for trozo in x['rango'].split('·'):
            blancos.add(trozo.replace(' ', '').replace('–', '-').strip())
    n = Counter()
    todas_hojas, rotulos = {}, {}
    for es, en, pes, pen in libros(carpeta):
        w = cargar(pen)
        corto_ = extraer_textos.corto(es)
        todas_hojas[corto_] = set(w.sheetnames)
        rotulos[corto_] = set(v.strip() for _, _, v in textos_libro(w) if len(v) <= 80) | dv_items(w)
    tcd = {x['es']: x['en'] for x in M['textos_cifra_derivada']}
    te = datos['te']
    donde_de = {c['es']: c['donde'] for c in te['cadenas']}
    for es, en, pes, pen in libros(carpeta):
        wb = cargar(pen)
        corto = extraer_textos.corto(es)
        etiquetas = set(v.strip() for _, _, v in textos_libro(wb))
        etiquetas |= dv_items(wb)
        for ws in wb.worksheets:
            for c in ws._cells.values():
                if c.data_type == 'f':
                    etiquetas.update(x.replace('""', '"') for x in re.findall(r'"((?:[^"]|"")*)"', c.value))
        otras = set().union(*[h for f, h in todas_hojas.items() if f != corto]) - set(wb.sheetnames)
        # (a) rangos D12 junto a su etiqueta
        for h, coord, t in textos_libro(wb):
            for m in RX_RANGO.finditer(t):
                n['rangos'] += 1
                rango = '%s-%s%%' % (m.group(1), m.group(2))
                antes = t[max(0, m.start() - 45):m.start()].lower()
                antes = re.split(r'[.;·|]\s|\(', antes)[-1]      # misma frase; «:» no corta
                                                                   # («for catering: 28-35%»)
                tipo = None
                for k, pats in ETIQUETAS_D12.items():
                    for p in pats:
                        for mm in re.finditer(p, antes):
                            if tipo is None or mm.end() > tipo[1]:
                                tipo = (k, mm.end())
                if not tipo:
                    continue
                lo, hi = float(m.group(1)), float(m.group(2))
                if (lo, hi) == d12[tipo[0]]:
                    n['rangos_d12_ok'] += 1
                elif rango in blancos:
                    n['rangos_lista_blanca'] += 1
                else:
                    R.fallo('%s:%s!%s: «%s» junto a %s ≠ D12 %s-%s%% · …%s…'
                            % (en, h, coord, rango, tipo[0], d12[tipo[0]][0], d12[tipo[0]][1],
                               t[max(0, m.start() - 50):m.end() + 5]))
        # (b) pestañas citadas (§2.2, D9/R2T-18). Comillas dobles = una pestaña o un rótulo
        #     visible de ESTE libro; o, si el texto nombra otro libro (su número, «BONUS» o «recipe
        #     cost card» = 01-08), una pestaña o un rótulo de ESE libro. Nunca un nombre que no
        #     exista en ninguno de ellos (p. ej. una pestaña renombrada).
        for h, coord, t in textos_citables(wb):
            nombrados = libros_nombrados(t) - {corto}
            for m in RX_COMILLAS_D.finditer(t):
                q = m.group(1)
                n['citas_dobles'] += 1
                if q in wb.sheetnames:
                    n['citas_pestana'] += 1
                    continue
                if cita_casa(q, [e for e in etiquetas if e != t.strip()]):
                    n['citas_rotulo'] += 1
                    continue
                if any(q in todas_hojas[b] or cita_casa(q, rotulos[b]) for b in nombrados):
                    n['citas_otro_libro'] += 1
                    continue
                if q in CITAS_BLANCAS:
                    n['citas_lista_blanca'] += 1
                    continue
                if q in otras:
                    R.fallo('%s:%s!%s: cita la pestaña «%s» de otro libro sin nombrar ese libro' % (en, h, coord, q))
                else:
                    R.fallo('%s:%s!%s: «"%s"» no es una pestaña ni un rótulo del libro (ni de los que nombra: %s)'
                            % (en, h, coord, q, sorted(nombrados)))
            for m in RX_COMILLAS_S.finditer(t):
                if m.group(1) in wb.sheetnames or m.group(1) in otras:
                    R.fallo('%s:%s!%s: pestaña «%s» entre comillas simples (D9: dobles rectas)'
                            % (en, h, coord, m.group(1)))
        # (c) cifras derivadas: la EN recalculada está donde estaba la ES
        for es_txt, en_txt in tcd.items():
            for d in donde_de.get(es_txt, []):
                if d['f'] != corto or d['t'] != 'celda':
                    continue
                ws = wb[mapas.hoja_en(d['h'])]
                v = ws[d['c']].value
                n['cifras_derivadas'] += 1
                if v != en_txt and not (d['h'] == 'Lista de precios'):
                    R.fallo('%s:%s!%s: cifra derivada no aplicada (%r)' % (en, ws.title, d['c'], (v or '')[:70]))
        # (d) Conversions!A2 y las celdas de tamaño citadas
        if 'Conversions' in wb.sheetnames:
            a2 = wb['Conversions']['A2'].value or ''
            if 'row %d' % datos['fin_conv'] not in a2:
                R.fallo('%s: Conversions!A2 no cita la fila %d: %r' % (en, datos['fin_conv'], a2))
            conv = {wb['Conversions'].cell(r, 1).value: r for r in range(5, 5 + datos['n'])}
            tam = {'bottle': 'B%d' % conv['bottle→ml'], 'can': 'B%d' % conv['can→ml']}
            for h, coord, t in textos_libro(wb):
                for m in re.finditer(r'cell (B\d+) \(row (bottle|can)→ml', t):
                    n['celdas_tamano_citadas'] += 1
                    if tam[m.group(2)] != m.group(1):
                        R.fallo('%s:%s!%s: cita %s para %s y es %s' % (en, h, coord, m.group(1), m.group(2), tam[m.group(2)]))
    # (f) D9: cabeceras de «Trim Loss Factors» y nota «yield = 100 % − trim loss»
    for es, en, pes, pen in libros(carpeta):
        wb = cargar(pen)
        if 'Trim Loss Factors' not in wb.sheetnames:
            continue
        ws = wb['Trim Loss Factors']
        cab = [ws.cell(4, c).value for c in range(1, 5)]
        if cab != ['Category', 'Typical trim loss %', 'Min trim loss %', 'Max trim loss %']:
            R.fallo('%s: cabeceras de Trim Loss Factors %r (D9)' % (en, cab))
        if 'Yield % = 100% − trim loss %' not in (ws['A2'].value or ''):
            R.fallo('%s: Trim Loss Factors!A2 sin «Yield % = 100% − trim loss %» (D9)' % en)
        if '12' not in (ws['A3'].value or '') or 'USDA Food Buying Guide' not in (ws['A3'].value or ''):
            R.fallo('%s: Trim Loss Factors!A3 no remite al 12 ni al USDA Food Buying Guide (D20)' % en)
        n['cabeceras_d9'] += 1
    # (e) calculadora 10 = D12
    pen = os.path.join(carpeta, mapas.FICHEROS['10-calculadora-pvp.xlsx'])
    ws = cargar(pen)['Menu Price Calculator']
    for f in M['calculadora_10']['filas']:
        r = f['fila']
        if (ws['B%d' % r].value, ws['C%d' % r].value, ws['D%d' % r].value) != (f['rotulo'], f['fc_min'], f['fc_max']):
            R.fallo('10 fila %d: %r ≠ D12 %r' % (r, (ws['B%d' % r].value, ws['C%d' % r].value, ws['D%d' % r].value),
                                                (f['rotulo'], f['fc_min'], f['fc_max'])))
    if ws['F18'].value != M['calculadora_10']['comision_delivery']:
        R.fallo('10: comisión de delivery %r' % ws['F18'].value)
    for k, v in sorted(n.items()):
        R.dato(k, v)


# ==========================================================================
# Gate 7 — censo-entregables
# ==========================================================================
def gate7(R, carpeta, datos):
    R.gate('7 censo-entregables')
    r = subprocess.run([sys.executable, CENSO_SCRIPT, '--only', carpeta.rstrip('/') + '/', '--letter',
                        '--fail', '--quiet'], capture_output=True, text=True)
    salida = (r.stdout + r.stderr).strip().splitlines()
    R.dato('salida', ' | '.join(salida[-3:]))
    if r.returncode != 0:
        R.fallo('censo-entregables.py --fail sale con %d: %s' % (r.returncode, ' | '.join(salida[-8:])))


# ==========================================================================
# Gate 7bis — libros 12 y 13
# ==========================================================================
def gate7bis(R, carpeta, datos):
    R.gate('7bis libros 12 y 13')
    M = datos['mercado']
    p12 = os.path.join(carpeta, mapas.FICHEROS['12-test-de-rendimiento.xlsx'])
    p01 = os.path.join(carpeta, mapas.FICHEROS['01-escandallo-estandar.xlsx'])
    p08 = os.path.join(carpeta, mapas.FICHEROS['08-food-truck.xlsx'])
    p13 = os.path.join(carpeta, mapas.FICHEROS['13-lista-precios-ingredientes.xlsx'])
    v12 = cargar(p12, data_only=True)
    d, c = v12['Butcher Yield Test'], v12['Cooking Loss Test']
    rend, factor, merma, coste = d['C32'].value, d['C34'].value, d['C35'].value, d['C37'].value
    R.dato('12 despiece', 'rendimiento %.4f · factor %.4f · merma %.4f · coste/porción %.4f'
           % (rend, factor, merma, coste))
    if not (0 < rend < 1):
        R.fallo('12: rendimiento %r fuera de (0, 1)' % rend)
    if not factor >= 1:
        R.fallo('12: factor de coste %r < 1' % factor)
    v01 = cargar(p01, data_only=True)['Recipe Cost Card']
    f01 = cargar(p01)['Recipe Cost Card']
    if abs(f01['H5'].value - round(merma, 4)) > 1e-9:
        R.fallo('01!H5 = %r ≠ merma del 12 %.4f' % (f01['H5'].value, merma))
    if f01['E5'].value != d['C36'].value or f01['C5'].value != 'lb' or f01['F5'].value != 'oz':
        R.fallo('01 fila 5 no es la porción del 12 (%r %r→%r)' % (f01['E5'].value, f01['C5'].value, f01['F5'].value))
    if abs(v01['J5'].value - coste) > TOL * coste:
        R.fallo('12: coste por porción %.4f ≠ 01!J5 %.4f' % (coste, v01['J5'].value))
    R.dato('01!J5', round(v01['J5'].value, 4))
    perd, cpc = c['C15'].value, c['C17'].value
    v08 = cargar(p08, data_only=True)['Smash Burger']
    if not (0 < perd < 1):
        R.fallo('12 cocción: pérdida %r' % perd)
    if abs(v08['J5'].value - cpc) > TOL * cpc:
        R.fallo('12 cocción: coste por porción cocinada %.4f ≠ 08!J5 %.4f' % (cpc, v08['J5'].value))
    R.dato('12 cocción', 'pérdida %.3f · coste %.4f · 08!J5 %.4f' % (perd, cpc, v08['J5'].value))
    # 13
    wf, wv = cargar(p13), cargar(p13, data_only=True)
    ws, wsv = wf['Price List'], wv['Price List']
    umbral = ws['B5'].value
    n_ing = n_alert = n_ok = 0
    for r in range(10, 155):
        if ws.cell(r, 1).value in (None, ''):
            continue
        n_ing += 1
        u = ws.cell(r, 6).value
        if u not in mapas.UNIDADES_EN:
            R.fallo('13!F%d: unidad base %r fuera de D7' % (r, u))
        j, k = wsv.cell(r, 10).value, wsv.cell(r, 11).value
        if _num(j):
            esperado = 'ALERT' if j > umbral else 'OK'
            if k != esperado:
                R.fallo('13!K%d = %r con variación %.4f (umbral %.2f)' % (r, k, j, umbral))
            n_alert += k == 'ALERT'
            n_ok += k == 'OK'
    if wsv['B4'].value != n_ing or wsv['B6'].value != n_alert:
        R.fallo('13: resumen B4/B6 = %r/%r ≠ %d/%d' % (wsv['B4'].value, wsv['B6'].value, n_ing, n_alert))
    if not (n_alert >= 1 and n_ok >= 1):
        R.fallo('13: la demo debe enseñar una ALERT y un OK (%d/%d)' % (n_alert, n_ok))
    R.dato('13', '%d ingredientes · %d ALERT · %d OK · umbral %s' % (n_ing, n_alert, n_ok, umbral))
    dv = [x for x in ws.data_validations.dataValidation if 'F10' in x.sqref]
    if not dv or dv[0].formula1 != mapas.DV_UNIDADES_EN:
        R.fallo('13: DV de unidad base ≠ lista D7')
    # pycel: la alerta salta con +10 % y no con +2 %
    xl = _pycel(p13)
    fila = next(r for r in range(10, 155) if _num(wsv.cell(r, 8).value) and wsv.cell(r, 9).value is None)
    h = wsv.cell(fila, 8).value
    _ev(xl, "'Price List'!K%d" % fila)                 # mete I y H de la fila en el mapa de pycel
    xl.set_value("'Price List'!I%d" % fila, h / 1.10)
    k1 = _ev(xl, "'Price List'!K%d" % fila)
    xl.set_value("'Price List'!I%d" % fila, h / 1.02)
    k2 = _ev(xl, "'Price List'!K%d" % fila)
    if (k1, k2) != ('ALERT', 'OK'):
        R.fallo('13 fila %d: +10 %% da %r y +2 %% da %r (se esperaba ALERT / OK)' % (fila, k1, k2))
    # 13 = fichas: precio por unidad base = precio de la ficha ÷ factor
    ing13 = {}
    for r in range(10, 155):
        a = ws.cell(r, 1).value
        if a:
            ing13[a] = (wsv.cell(r, 8).value, ws.cell(r, 6).value, r)
    nombre13 = {f['id']: f['ingrediente'] for f in M['libro13']['filas']}
    factores = {x['clave']: x['factor'] for x in datos['conv']}
    comprobadas = 0
    for rec in M['recetas']:
        fname = [f for f in mapas.LIBROS_ES if f.startswith(rec['libro'])][0]
        pen = os.path.join(carpeta, mapas.FICHEROS[fname])
        wvv = datos['cache_v'].setdefault(pen, cargar(pen, data_only=True))
        wff = datos['cache_f'].setdefault(pen, cargar(pen))
        for f in rec['filas']:
            nom = nombre13.get(f['id'])
            if nom is None or nom not in ing13:
                continue
            precio_ficha = wvv[rec['hoja_en']].cell(f['fila'], 4).value
            ud = wff[rec['hoja_en']].cell(f['fila'], 3).value
            h13, base, r13 = ing13[nom]
            fac = 1.0 if ud == base else factores.get(ud + '→' + base)
            if fac is None:
                R.fallo('13 %r: sin factor %s→%s' % (nom, ud, base))
                continue
            esperado = precio_ficha / fac
            comprobadas += 1
            if abs(esperado - h13) > max(TOL * h13, 0.00015):
                R.fallo('13!H%d %r: %.4f ≠ ficha %s fila %d %.4f/%s ÷ %.6f = %.4f'
                        % (r13, nom, h13, rec['hoja_en'], f['fila'], precio_ficha, ud, fac, esperado))
    R.dato('13 = fichas', '%d filas de ficha comprobadas' % comprobadas)


# ==========================================================================
# Gate PDF — el bono contra los xlsx (T11)
# ==========================================================================
RX_XLSX = re.compile(r'\b(?:\d{2}|BONUS)-[a-z0-9-]+\.xlsx\b')
RX_TAB = re.compile(r'"([^"\n]{1,40})"(?:\s+and\s+"([^"\n]{1,40})")?\s+tabs?\b')
RX_RANGO_PCT = re.compile(r'\b(\d{1,2})-(\d{1,2})%')


def _pct(x):
    return int(round(x * 100))


def gate_pdf(R, carpeta, datos):
    R.gate('pdf bono')
    from pypdf import PdfReader
    path = os.path.join(carpeta, aplicar_en.PDF_EN)
    if not os.path.isfile(path):
        R.fallo('no existe %s' % path)
        return
    with open(os.devnull, 'w') as dn, contextlib.redirect_stderr(dn):
        rd = PdfReader(path)
        paginas = [p.extract_text() or '' for p in rd.pages]
    texto = '\n'.join(paginas)
    R.dato('paginas', len(rd.pages))
    if len(rd.pages) != PAGINAS_PDF:
        R.fallo('%d páginas ≠ %d (la cifra que dirá la landing EN)' % (len(rd.pages), PAGINAS_PDF))
    for i, p in enumerate(rd.pages):
        mb = p.mediabox
        if (round(float(mb.width)), round(float(mb.height))) != (612, 792):
            R.fallo('página %d: %sx%s pt, no US Letter' % (i + 1, mb.width, mb.height))
    # restos y caracteres
    malos = no_permitidos(texto)
    if malos:
        R.fallo('carácter fuera de la lista blanca en el PDF: %r' % ''.join(malos))
    for motivo, ctx in restos_espanol(texto):
        R.fallo('PDF: %s · …%s…' % (motivo, ctx))
    m = RX_ID_INTERNO.search(texto)
    if m:
        R.fallo('PDF: ID interno %r' % m.group(0))
    # ficheros y pestañas citados
    ficheros_en = set(mapas.FICHEROS.values())
    hojas_en = set()
    libros_en = {}
    for es, en, pes, pen in libros(carpeta):
        w = datos['cache_f'].setdefault(pen, cargar(pen))
        hojas_en |= set(w.sheetnames)
        libros_en[extraer_textos.corto(es)] = w
    citados = set(RX_XLSX.findall(texto))
    R.dato('ficheros citados', len(citados))
    for f in sorted(citados - ficheros_en):
        R.fallo('PDF cita un fichero que no existe: %s' % f)
    n_tabs = 0
    for m in RX_TAB.finditer(texto):
        for q in (m.group(1), m.group(2)):
            if q:
                n_tabs += 1
                if q not in hojas_en:
                    R.fallo('PDF cita la pestaña «%s», que no existe en ningún xlsx' % q)
    R.dato('pestañas citadas', n_tabs)
    # tabla de trim loss = «Trim Loss Factors» del 01
    ws = libros_en['01']['Trim Loss Factors']
    ref = {ws.cell(r, 1).value: (ws.cell(r, 2).value, ws.cell(r, 3).value, ws.cell(r, 4).value)
           for r in range(5, 26) if ws.cell(r, 1).value}
    i0 = texto.find('Reference trim loss by category')
    filas = 0
    for cat, (tip, lo, hi) in ref.items():
        mm = re.search(r'(?m)^%s\n(\d+)%%\n(\d+)-(\d+)%%$' % re.escape(cat), texto[i0:i0 + 3000]) if i0 >= 0 else None
        if not mm:
            continue
        filas += 1
        if tuple(int(x) for x in mm.groups()) != (_pct(tip), _pct(lo), _pct(hi)):
            R.fallo('PDF, trim loss de «%s»: %s%% (%s-%s%%) ≠ xlsx %d%% (%d-%d%%)'
                    % ((cat,) + mm.groups() + (_pct(tip), _pct(lo), _pct(hi))))
    R.dato('filas trim loss', filas)
    if filas < 8:
        R.fallo('PDF: solo %d filas de la tabla de trim loss casan con «Trim Loss Factors» (se esperan 8)' % filas)
    # tabla D12 = calculadora 10 (B9:D18)
    wc = libros_en['10']['Menu Price Calculator']
    j0 = texto.find('Venue type')
    for r in range(9, 19):
        rot, lo, hi = wc['B%d' % r].value, wc['C%d' % r].value, wc['D%d' % r].value
        k = texto.find('\n' + rot, j0) if j0 >= 0 else -1
        mm = RX_RANGO_PCT.search(texto, k) if k >= 0 else None
        if not mm or mm.start() - k > 160:
            R.fallo('PDF: la tabla D12 no tiene la fila «%s» con su rango' % rot)
        elif (int(mm.group(1)), int(mm.group(2))) != (_pct(lo), _pct(hi)):
            R.fallo('PDF, D12 «%s»: %s-%s%% ≠ calculadora 10 %d-%d%%' % (rot, mm.group(1), mm.group(2),
                                                                         _pct(lo), _pct(hi)))


# ==========================================================================
# Autotest (defectos inyectados en una copia)
# ==========================================================================
def _mutar(carpeta, fname_en, fn, modo='wb'):
    p = os.path.join(carpeta, fname_en)
    if modo == 'wb':                      # openpyxl + inject_cache
        wb = openpyxl.load_workbook(p)
        fn(wb)
        wb.save(p)
        subprocess.run([sys.executable, '-W', 'ignore', INJECT, p], capture_output=True, text=True)
    else:                                 # 'zip' / 'pdf': la función recibe la ruta y edita el fichero tal cual
        fn(p)


def _sin_v(path, hoja_xml, celdas):
    """Quita el <v> de unas celdas en el XML (T5): simula un libro guardado sin caché."""
    with zipfile.ZipFile(path) as z:
        infos = z.infolist()
        partes = {i.filename: z.read(i.filename) for i in infos}
    x = partes[hoja_xml].decode('utf-8')
    for c in celdas:
        x, k = re.subn(r'(<c r="%s"[^>]*>(?:(?!</c>).)*?)<v>[^<]*</v>' % c, r'\1', x, flags=re.S)
        assert k == 1, c
    partes[hoja_xml] = x.encode('utf-8')
    tmp = path + '.tmp'
    with zipfile.ZipFile(tmp, 'w', zipfile.ZIP_DEFLATED) as zout:
        for i in infos:
            zout.writestr(i, partes[i.filename])
    os.replace(tmp, path)


def _parte_hoja(path, nombre):
    with zipfile.ZipFile(path) as z:
        wbx = z.read('xl/workbook.xml').decode('utf-8')
        rels = z.read('xl/_rels/workbook.xml.rels').decode('utf-8')
    rid = re.search(r'<sheet\b[^>]*name="%s"[^>]*r:id="([^"]+)"' % re.escape(nombre), wbx).group(1)
    tag = [t for t in re.findall(r'<Relationship\b[^>]*>', rels) if 'Id="%s"' % rid in t][0]
    t = re.search(r'Target="([^"]+)"', tag).group(1)
    return t.lstrip('/') if t.startswith('/') else 'xl/' + t


def _pdf_sin_ultima(path):
    from pypdf import PdfReader, PdfWriter
    with open(os.devnull, 'w') as dn, contextlib.redirect_stderr(dn):
        rd, wr = PdfReader(path), PdfWriter()
        for pg in rd.pages[:-1]:
            wr.add_page(pg)
        with open(path, 'wb') as fh:
            wr.write(fh)


def _cf_literal(ws, de, a):
    for cf in ws.conditional_formatting:
        for regla in cf.rules:
            if regla.formula and ('"%s"' % de) in regla.formula[0]:
                regla.formula = [x.replace('"%s"' % de, '"%s"' % a) for x in regla.formula]


P01 = '01-standard-recipe-cost-card.xlsx'
# (gate, descripción, fichero, mutación, subcadena que TIENE que salir en un fallo, modo)
DEFECTOS = [
    ('1', 'hoja citada inexistente', '02-tasting-menu.xlsx',
     lambda wb: setattr(wb['Summary']['C5'], 'value', "=IF('1. Aperitivo'!$J$21=0,\"\",'1. Amuse-Bouche'!$J$25)"),
     'hoja inexistente', 'wb'),
    ('1', 'DV de unidades sin sustituir', P01,
     lambda wb: setattr([d for d in wb['Recipe Cost Card'].data_validations.dataValidation
                         if d.formula1 == mapas.DV_UNIDADES_EN][0], 'formula1', aplicar_en.UNIDADES_ES_DV),
     'no es la lista D7', 'wb'),
    ('1', 'T1: literal de fórmula en español («revisa unidades»)', P01,
     lambda wb: setattr(wb['Recipe Cost Card']['J5'], 'value',
                        wb['Recipe Cost Card']['J5'].value.replace('check units', 'revisa unidades')),
     'fórmula EN', 'wb'),
    ('1', 'T1: CF que busca «ALERTA» (13)', '13-ingredient-price-list.xlsx',
     lambda wb: _cf_literal(wb['Price List'], 'ALERT', 'ALERTA'), 'formato condicional distinto', 'wb'),
    ('2', 'T1: literal de CF en español (09)', '09-food-waste-tracker.xlsx',
     lambda wb: _cf_literal(wb['Weekly Waste Log'], 'ALERT', 'ALERTA'), 'oficio ES', 'wb'),
    ('2', 'resto de español en un rótulo', '10-menu-price-calculator.xlsx',
     lambda wb: setattr(wb['Menu Price Calculator']['B2'], 'value', 'Calculadora de PVP Sugerido'), ' ES «', 'wb'),
    ('2', '«€» en un formato', '11-monthly-food-cost-dashboard.xlsx',
     lambda wb: setattr(wb['Dashboard']['C7'], 'number_format', '#,##0 €'), '«€»', 'wb'),
    ('2', 'CJK inyectado', '06-catering-and-events.xlsx',
     lambda wb: setattr(wb['Event Checklist']['A1'], 'value', 'Event Checklist 鬼笔鹅膏菌 — Catering'),
     'lista blanca', 'wb'),
    ('2', 'T6: paréntesis de ancho completo', '06-catering-and-events.xlsx',
     lambda wb: setattr(wb['Client Proposal']['D7'], 'value', 'Price per guest（incl. tax）'), 'lista blanca', 'wb'),
    ('2', 'T2: ID interno de la SPEC en un texto', '11-monthly-food-cost-dashboard.xlsx',
     lambda wb: setattr(wb['Instructions']['B47'], 'value', wb['Instructions']['B47'].value + ' (D19)'),
     'ID interno', 'wb'),
    ('3', 'clave duplicada (mayúsculas) en Conversions', '05-pastry-and-bakery.xlsx',
     lambda wb: (setattr(wb['Conversions']['A%d' % (mapas.FILA0 + mapas.n_claves())], 'value', 'LB→lb'),
                 setattr(wb['Conversions']['B%d' % (mapas.FILA0 + mapas.n_claves())], 'value', 1)),
     'claves duplicadas', 'wb'),
    ('3', 'T5: fórmulas sin caché (<v> quitado)', P01,
     lambda p: _sin_v(p, _parte_hoja(p, 'Recipe Cost Card'), ['J33', 'J35', 'J36', 'J37']), 'sin caché', 'zip'),
    ('3', 'T7: fila de envase con el IFERROR viejo (0 con la celda vacía)', P01,
     lambda wb: setattr(wb['Conversions']['B148'], 'value', '=IFERROR(B147/29.5735295625,"?")'),
     'celda de tamaño vacía', 'wb'),
    ('4', 'cantidad que saca el food cost de D12', '07-cafe-and-brunch.xlsx',
     lambda wb: setattr(wb['Avocado Toast']['E5'], 'value', wb['Avocado Toast']['E5'].value * 4), 'fuera de', 'wb'),
    ('4', 'CHEF-03: pase de menú con precio propio', '02-tasting-menu.xlsx',
     lambda wb: setattr(wb['2. Starter']['I34'], 'value', 50.91), 'precio propio', 'wb'),
    ('5', 'hoja en A4', '09-food-waste-tracker.xlsx',
     lambda wb: setattr(wb['Trend'].page_setup, 'paperSize', 9), 'Letter', 'wb'),
    ('5', 'fecha en formato propio', 'BONUS-inventory-and-waste-control.xlsx',
     lambda wb: setattr(wb['Waste Checklist']['A4'], 'number_format', 'd/m/yyyy'), 'numFmtId', 'wb'),
    ('5', 'línea de versión ausente', '03-prix-fixe-lunch-menu.xlsx',
     lambda wb: [setattr(c, 'value', None) for c in list(wb['Instructions']._cells.values())
                 if isinstance(c.value, str) and c.value.startswith('Version ')], 'línea de versión', 'wb'),
    ('5', 'CHEF-01: impuesto del 10 % en una pestaña vacía', '02-tasting-menu.xlsx',
     lambda wb: setattr(wb['6. Course']['I29'], 'value', 0.1), 'Tax rate', 'wb'),
    ('6', 'pestaña entre comillas simples', '08-food-truck.xlsx',
     lambda wb: setattr(wb['Instructions']['B8'], 'value', "▸ Start on the 'Smash Burger' tab."),
     'comillas simples', 'wb'),
    ('6', 'rango D12 incoherente', '10-menu-price-calculator.xlsx',
     lambda wb: setattr(wb['Instructions']['B8'], 'value', '▸ Fine dining runs at 25-28% food cost.'), '≠ D12', 'wb'),
    ('6', 'T4: pestaña citada que no existe (prefijo de un rótulo)', P01,
     lambda wb: setattr(wb['Instructions']['B11'], 'value',
                        '▸ Open the "Conversion" tab and the "Trim Loss" tab.'), 'no es una pestaña', 'wb'),
    ('7bis', '13 descuadrado con la ficha', '13-ingredient-price-list.xlsx',
     lambda wb: setattr(wb['Price List']['G11'], 'value', wb['Price List']['G11'].value * 1.2), '≠ ficha', 'wb'),
    ('7bis', 'porción del 12 distinta de la 01', '12-yield-test.xlsx',
     lambda wb: setattr(wb['Butcher Yield Test']['C9'], 'value', 19.5), '01!J5', 'wb'),
    ('pdf', 'PDF con una página menos', aplicar_en.PDF_EN, _pdf_sin_ultima, 'páginas', 'pdf'),
]


def autotest(carpeta, mes, datos):
    """T11: 1) cada gate usado pasa limpio sobre la copia SIN mutar (si no, el autotest no prueba nada);
    2) cada defecto tiene que producir un fallo que contenga SU subcadena (otro fallo cualquiera no cuenta)."""
    base = os.path.join(os.path.dirname(os.path.abspath(carpeta)), 'rck-autotest')
    gates = sorted({d[0] for d in DEFECTOS}, key=list(GATES).index)
    R0 = ejecutar(carpeta, mes, gates, datos, silencioso=True)
    limpios = True
    for g in R0.gates.values():
        if g['fallos']:
            limpios = False
            print('  [!!] la copia limpia ya falla en «%s»: %s' % (list(R0.gates)[list(R0.gates.values()).index(g)],
                                                               g['fallos'][0][:100]))
    print('  base limpia: %s (gates %s)' % ('OK' if limpios else 'NO', ', '.join(gates)))
    detectados = []
    for gate, desc, fname, fn, esperado, modo in DEFECTOS:
        aplicar_en.termica()
        if os.path.isdir(base):
            shutil.rmtree(base)
        shutil.copytree(carpeta, base)
        _mutar(base, fname, fn, modo)
        R = ejecutar(base, mes, [gate], datos, silencioso=True)
        fallos = R.gates[list(R.gates)[0]]['fallos']
        propios = [f for f in fallos if esperado in f]
        ok = bool(propios)
        detectados.append(ok)
        print('  [%s] gate %-4s %-58s → %s' % ('OK' if ok else '!!', gate, desc[:58],
                                                (propios[0][:90] if propios else 'NO DETECTADO (%d fallos ajenos: %s)'
                                                 % (len(fallos), fallos[0][:60] if fallos else '—'))))
    shutil.rmtree(base, ignore_errors=True)
    print('autotest: %d/%d defectos detectados con su mensaje · base limpia: %s'
          % (sum(detectados), len(detectados), 'sí' if limpios else 'NO'))
    return all(detectados) and limpios


# ==========================================================================
# main
# ==========================================================================
GATES = OrderedDict([('1', gate1), ('2', gate2), ('3', gate3), ('4', gate4), ('5', gate5),
                     ('6', gate6), ('7', gate7), ('7bis', gate7bis), ('pdf', gate_pdf)])


def cargar_datos():
    d = aplicar_en.cargar_datos()
    d['cache_v'], d['cache_f'] = {}, {}
    return d


def ejecutar(carpeta, mes, solo, datos, silencioso=False):
    R = Resultado()
    datos['cache_v'], datos['cache_f'] = {}, {}
    for k, fn in GATES.items():
        if solo and k not in solo:
            continue
        aplicar_en.termica()
        if k == '5':
            fn(R, carpeta, datos, mes)
        else:
            fn(R, carpeta, datos)
        if not silencioso:
            g = R.gates[R.actual]
            print('%s %-24s %s' % ('✔' if not g['fallos'] else '✘', R.actual,
                                  'OK' if not g['fallos'] else '%d fallos' % len(g['fallos'])), flush=True)
    return R


def main():
    ap = argparse.ArgumentParser(description='Gates F2 del Recipe Costing Kit Pro (EN)')
    ap.add_argument('--dir', default=None, help='carpeta con los 14 xlsx EN (por defecto, la del dry-run)')
    ap.add_argument('--mes', default=None)
    ap.add_argument('--solo', default=None, help='gates separados por comas (1,2,3,4,5,6,7,7bis,pdf)')
    ap.add_argument('--json', default=None)
    ap.add_argument('--verbose', action='store_true')
    ap.add_argument('--autotest', action='store_true')
    args = ap.parse_args()
    carpeta = aplicar_en.carpeta_dryrun(args.dir)
    mes = args.mes or aplicar_en.mes_por_defecto()
    datos = cargar_datos()
    if args.autotest:
        sys.exit(0 if autotest(carpeta, mes, datos) else 1)
    solo = args.solo.split(',') if args.solo else None
    R = ejecutar(carpeta, mes, solo, datos)
    total = 0
    for nombre, g in R.gates.items():
        total += len(g['fallos'])
        if g['fallos'] or args.verbose:
            print('\n== %s' % nombre)
            for k, v in g['datos'].items():
                print('   · %s: %s' % (k, v))
            for f in g['fallos'][:60 if not args.verbose else None]:
                print('   ✘ ' + f)
            if len(g['fallos']) > 60 and not args.verbose:
                print('   … y %d más' % (len(g['fallos']) - 60))
    if args.json:
        with open(args.json, 'w', encoding='utf-8') as fh:
            json.dump(R.gates, fh, ensure_ascii=False, indent=1, default=str)
    print('\n%s: %d fallos en %d gates' % ('OK' if not total else 'FALLA', total, len(R.gates)))
    sys.exit(1 if total else 0)


if __name__ == '__main__':
    main()
