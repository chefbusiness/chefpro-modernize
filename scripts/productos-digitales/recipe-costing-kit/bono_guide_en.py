#!/usr/bin/env python3
"""
bono_guide_en.py — Maqueta el bono EN «How to Reduce Food Cost in 30 Days»
(Food Cost Kit Pro, SPEC §2.6 y D9 bis) desde Markdown a PDF (US Letter) y, si se
pide, a DOCX de trabajo.

Copia PARAMETRIZADA de `kit-escandallos-v2_0/bono_guia.py` (el maquetador del
bono ES, que no se toca): mismo parser Markdown, mismo saneado WinAnsi y mismo
diseño (dorado #B8860B, Helvetica, tablas con cabecera dorada). Lo que cambia
está en PARAMS: cabecera, pie «Page {n}», título, marcador de la línea de
metadatos, papel Letter y metadatos del PDF en inglés.

    /usr/local/bin/python3 bono_guide_en.py <entrada.md> <salida_pdf_sin_extension>
                                            [--docx <salida_docx_sin_extension>]
    /usr/local/bin/python3 bono_guide_en.py <entrada.md> --solo-gates
    /usr/local/bin/python3 bono_guide_en.py --autotest

Gates (abortan con código 3; el ES solo tenía el primero):
  1. GLIFOS: tras el saneado, cero caracteres fuera de WinAnsi (cp1252): lo que
     Helvetica no sabe pintar sale como un cuadrado negro. También se revisan
     los textos fijos (cabecera, pie, título, marcador).
  2. RESTOS DE ESPAÑOL: cero signos (¿ ¡ « » € º ª), cero palabras con
     tilde o eñe fuera de la lista blanca (café, purée…), cero palabras
     funcionales o de oficio en español (de, la, que, IVA, escandallo, merma…)
     y cero «entrée» (SPEC D8). NO exige ñ/€/«» como el ES: aquí son fallo.
  3. Tras maquetar: se relee el PDF con pypdf, se cuentan las páginas, se
     comprueba el pie «Page n» en cada una y se repite el gate 2 sobre el texto
     extraído (lo que el cliente ve de verdad).

El .docx NO es entregable (COM-B18 del ES): solo se escribe con --docx, fuera
de la carpeta de entrega. No toca `astro-site/public/dl/`: escribe donde se le
diga.
"""
import argparse
import re
import sys

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, Cm, RGBColor

from reportlab.lib.pagesizes import letter
from reportlab.lib import colors
from reportlab.lib.units import cm
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.platypus import (SimpleDocTemplate, Paragraph, Spacer, Table,
                                TableStyle, ListFlowable, ListItem,
                                Preformatted)

# ==========================================================================
# Parámetros (lo único que diferencia este maquetador del ES)
# ==========================================================================
PARAMS = {
    'pagesize': letter,                                   # SPEC D13: US Letter
    'cabecera': 'AI Chef Pro · Food Cost Kit Pro',
    'pie': 'Page {n}',
    'titulo_pdf': 'How to Reduce Food Cost in 30 Days',       # SPEC D9 bis (= mapas.TITULOS)
    # La línea en negrita bajo el H1 que se pinta como subtítulo (y no se
    # repite en el cuerpo). Se localiza por este marcador.
    'marcador_meta': 'A bonus guide from Food Cost Kit Pro',
    'meta_por_defecto': 'A bonus guide from Food Cost Kit Pro by AI Chef Pro',
    'autor_pdf': 'AI Chef Pro',
    'asunto_pdf': 'Food Cost Kit Pro · Bonus guide',
    'creador_pdf': 'AI Chef Pro · aichef.pro',
    'keywords_pdf': 'food cost, how to reduce food cost, menu engineering, prime cost, '
                    'Food Cost Kit Pro, AI Chef Pro',
}

GOLD = '#B8860B'
DARK = '#1A1A1A'

# ==========================================================================
# Gate 1 — Saneado de caracteres (reportlab + Helvetica = WinAnsi/cp1252)
# ==========================================================================
# Mismo mapa que el ES (COM-A01): sin él, ✅ y ❌ colapsan en el mismo ■ y las
# casillas ☐ del checklist parecen ya marcadas.
MAPA_WINANSI = {
    '☐': '[  ]',      # ☐ casilla vacía
    '☑': '[X]',       # ☑
    '☒': '[X]',       # ☒
    '✅': 'OK',        # ✅
    '✔': 'OK',        # ✔
    '❌': 'X',         # ❌
    '✖': 'X',         # ✖
    '⭐': '*',         # ⭐
    '★': '*',         # ★
    # Van seguidos de su propia palabra en la tabla de menu engineering
    # («🐴 Plowhorse»): se quitan y la palabra se queda sola.
    '\U0001F434': '',      # 🐴 Plowhorse
    '\U0001F9E9': '',      # 🧩 Puzzle
    '\U0001F436': '',      # 🐶 Dog
    '→': '->',        # →
    '←': '<-',        # ←
    '─': '-',         # ─
    '━': '-',
    '−': '-',         # − (menos matemático)
    ' ': ' ',         # espacio duro
    ' ': ' ',         # espacio fino
    '‑': '-',         # guion no separable
}


def sanear(texto):
    """Sustituye lo que Helvetica no sabe pintar. Lo que quede fuera de cp1252
    y no esté en el mapa se elimina (el gate lo cuenta ANTES y aborta)."""
    if not isinstance(texto, str):
        return texto
    fuera = []
    for ch in texto:
        if ch in MAPA_WINANSI:
            fuera.append(MAPA_WINANSI[ch])
            continue
        try:
            ch.encode('cp1252')
        except UnicodeEncodeError:
            continue
        fuera.append(ch)
    return ''.join(fuera)


def sanear_bloques(bloques):
    fuera = []
    for tipo, contenido in bloques:
        if isinstance(contenido, str):
            contenido = sanear(contenido)
        elif tipo == 'table':
            contenido = [[re.sub(r'\s{2,}', ' ', sanear(c)).strip()
                          for c in fila] for fila in contenido]
        elif isinstance(contenido, list):
            contenido = [sanear(x) for x in contenido]
        fuera.append((tipo, contenido))
    return fuera


def _trozos(bloques):
    for _, contenido in bloques:
        if isinstance(contenido, str):
            yield contenido
        elif isinstance(contenido, list) and contenido and \
                isinstance(contenido[0], list):
            for fila in contenido:
                for c in fila:
                    yield c
        elif isinstance(contenido, list):
            for x in contenido:
                if isinstance(x, str):
                    yield x


def restos_no_winansi(textos):
    """Caracteres fuera de cp1252 que NO cubre MAPA_WINANSI: saldrían como ■
    o desaparecerían. Recibe un iterable de cadenas."""
    malos = {}
    for t in textos:
        for ch in t:
            if ch in MAPA_WINANSI:
                continue
            try:
                ch.encode('cp1252')
            except UnicodeEncodeError:
                malos[ch] = malos.get(ch, 0) + 1
    return malos


def fuera_de_cp1252(textos):
    """Todo lo que no es cp1252, cubierto o no por el mapa (para el informe)."""
    n = 0
    for t in textos:
        for ch in t:
            try:
                ch.encode('cp1252')
            except UnicodeEncodeError:
                n += 1
    return n


# ==========================================================================
# Gate 2 — Restos de español
# ==========================================================================
SIGNOS_ES = set('¿¡«»€ºª')
RX_TILDE = re.compile(r'[áéíóúüñÁÉÍÓÚÜÑ]')
# Palabras con tilde que el inglés sí usa (SPEC §4 gate 2: açaí, jalapeño,
# Ibérico… más las de cocina que salen en esta guía).
LISTA_BLANCA_TILDE = {
    'café', 'cafés', 'purée', 'purées', 'puréed', 'sauté', 'sautés',
    'sautéed', 'açaí', 'jalapeño', 'jalapeños', 'crème', 'brûlée',
    'ibérico', 'flambé', 'soufflé', 'consommé', 'fumé',
}
# Funcionales del español que no son palabras inglesas (sin «son», «sin»,
# «hay», «a», «lo»: chocan con el inglés).
FUNCIONALES_ES = {
    'de', 'del', 'la', 'las', 'el', 'los', 'y', 'que', 'con', 'para', 'por',
    'una', 'unos', 'unas', 'es', 'en', 'al', 'se', 'su', 'sus', 'como', 'pero',
    'muy', 'cuando', 'donde', 'porque', 'este', 'esta', 'esto', 'estos',
    'desde', 'hasta', 'sobre', 'entre', 'tu', 'tus', 'te', 'mi', 'mis',
}
# Oficio y restos típicos del .md ES.
OFICIO_ES = {
    'iva', 'escandallo', 'escandallos', 'escandallar', 'merma', 'mermas',
    'albaran', 'albaranes', 'proveedor', 'proveedores', 'plato', 'platos',
    'cocina', 'cocinero', 'ventas', 'compras', 'semana', 'semanas', 'coste',
    'costes', 'racion', 'raciones', 'pvp', 'tpv', 'euro', 'euros', 'bono',
    'guia', 'pagina', 'paginas', 'pestana', 'plantilla', 'plantillas',
    'pulpo', 'lubina', 'corvina', 'carrillera', 'solomillo', 'lomo',
    'restaurante', 'consultor', 'gastronomico', 'anos', 'caballo', 'perro',
    'estrella', 'rotura', 'inventario', 'factura', 'ingenieria',
}
PROHIBIDAS_EN = {'entrée', 'entree'}      # SPEC D8: nunca «entrée»

RX_URL = re.compile(r'https?://\S+|\b[\w-]+(?:\.[\w-]+)*\.(?:es|com|pro|gov|org|uk|co)\b[/\w.#-]*',
                    re.I)
RX_PALABRA = re.compile(r"[A-Za-zÀ-ÖØ-öø-ÿ]+")


def _sin_tilde(p):
    return (p.replace('á', 'a').replace('é', 'e').replace('í', 'i')
             .replace('ó', 'o').replace('ú', 'u').replace('ü', 'u')
             .replace('ñ', 'n'))


def restos_espanol(textos):
    """Devuelve [(motivo, fragmento)] con cada resto de español encontrado."""
    hallazgos = []
    for t in textos:
        if not isinstance(t, str):
            continue
        limpio = RX_URL.sub(' ', t)
        for ch in limpio:
            if ch in SIGNOS_ES:
                hallazgos.append(('signo', ch))
        for m in RX_PALABRA.finditer(limpio):
            w = m.group(0)
            wl = w.lower()
            ctx = limpio[max(0, m.start() - 25):m.end() + 25].replace('\n', ' ')
            if wl in PROHIBIDAS_EN:
                hallazgos.append(('prohibida D8', ctx))
            elif RX_TILDE.search(w) and wl not in LISTA_BLANCA_TILDE:
                hallazgos.append(('tilde/eñe', ctx))
            elif wl in FUNCIONALES_ES:
                hallazgos.append(('funcional ES «%s»' % wl, ctx))
            elif _sin_tilde(wl) in OFICIO_ES and wl not in LISTA_BLANCA_TILDE:
                hallazgos.append(('oficio ES «%s»' % wl, ctx))
    return hallazgos


# ==========================================================================
# Parser Markdown → bloques (idéntico al ES)
# ==========================================================================
def parsear(md_text):
    """Lista de bloques tipados: h1, h2, h3, p, ul, ol, table, code, hr."""
    lineas = md_text.split('\n')
    bloques = []
    i, n = 0, len(lineas)
    while i < n:
        linea = lineas[i].rstrip()

        if not linea.strip():
            i += 1
            continue

        if linea.strip() == '---':
            bloques.append(('hr', None))
            i += 1
            continue

        if linea.startswith('# '):
            bloques.append(('h1', linea[2:].strip()))
            i += 1
            continue
        if linea.startswith('## '):
            bloques.append(('h2', linea[3:].strip()))
            i += 1
            continue
        if linea.startswith('### '):
            bloques.append(('h3', linea[4:].strip()))
            i += 1
            continue

        if linea.strip().startswith('```'):
            codigo = []
            i += 1
            while i < n and not lineas[i].strip().startswith('```'):
                codigo.append(lineas[i])
                i += 1
            i += 1
            bloques.append(('code', '\n'.join(codigo)))
            continue

        if linea.strip().startswith('|') and i + 1 < n and \
                re.match(r'^\s*\|?[\s:|-]+\|?\s*$', lineas[i + 1]) and \
                '-' in lineas[i + 1]:
            filas = [linea]
            i += 2
            while i < n and lineas[i].strip().startswith('|'):
                filas.append(lineas[i])
                i += 1
            tabla = []
            for f in filas:
                celdas = [c.strip() for c in f.strip().strip('|').split('|')]
                tabla.append(celdas)
            bloques.append(('table', tabla))
            continue

        if re.match(r'^\s*[-*]\s+', linea):
            items = []
            while i < n and re.match(r'^\s*[-*]\s+', lineas[i]):
                items.append(re.sub(r'^\s*[-*]\s+', '', lineas[i].rstrip()))
                i += 1
            bloques.append(('ul', items))
            continue

        if re.match(r'^\s*\d+\.\s+', linea):
            items = []
            while i < n and re.match(r'^\s*\d+\.\s+', lineas[i]):
                items.append(re.sub(r'^\s*\d+\.\s+', '', lineas[i].rstrip()))
                i += 1
            bloques.append(('ol', items))
            continue

        parrafo = [linea]
        i += 1
        while i < n and lineas[i].strip() and \
                not lineas[i].startswith('#') and \
                lineas[i].strip() != '---' and \
                not lineas[i].strip().startswith('```') and \
                not lineas[i].strip().startswith('|') and \
                not re.match(r'^\s*[-*]\s+', lineas[i]) and \
                not re.match(r'^\s*\d+\.\s+', lineas[i]):
            parrafo.append(lineas[i].rstrip())
            i += 1
        bloques.append(('p', ' '.join(parrafo)))

    return bloques


RX_CURSIVA_PARRAFO = re.compile(r'^\*(?!\*)(.+?)(?<!\*)\*$', re.S)


def cursiva_entera(texto):
    """COM-M11 del ES: la firma (última línea, en cursiva) sin asteriscos."""
    if not isinstance(texto, str):
        return texto, False
    m = RX_CURSIVA_PARRAFO.match(texto.strip())
    if m and '*' not in m.group(1):
        return m.group(1).strip(), True
    return texto, False


def negrita_runs(texto):
    partes = re.split(r'(\*\*[^*]+\*\*)', texto)
    runs = []
    for p in partes:
        if not p:
            continue
        if p.startswith('**') and p.endswith('**'):
            runs.append((p[2:-2], True))
        else:
            runs.append((p, False))
    return runs


# ==========================================================================
# DOCX (solo copia de trabajo, --docx)
# ==========================================================================
def construir_docx(bloques, salida_docx, titulo_meta):
    doc = Document()
    estilo_normal = doc.styles['Normal']
    estilo_normal.font.name = 'Calibri'
    estilo_normal.font.size = Pt(11)
    for section in doc.sections:
        section.page_width = Cm(21.59)                    # Letter
        section.page_height = Cm(27.94)
        section.top_margin = Cm(2.2)
        section.bottom_margin = Cm(2.2)
        section.left_margin = Cm(2.2)
        section.right_margin = Cm(2.2)

    def escribir_runs(p, texto, negrita_base=False):
        for texto_run, es_negrita in negrita_runs(texto):
            r = p.add_run(texto_run)
            r.bold = es_negrita or negrita_base

    primero_h1 = True
    for tipo, contenido in bloques:
        if tipo == 'h1':
            if primero_h1:
                h = doc.add_heading(contenido, level=0)
                h.alignment = WD_ALIGN_PARAGRAPH.CENTER
                sub = doc.add_paragraph(titulo_meta)
                sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
                sub.runs[0].italic = True
                sub.runs[0].font.color.rgb = RGBColor(0x66, 0x66, 0x66)
                doc.add_paragraph()
                primero_h1 = False
            else:
                doc.add_heading(contenido, level=1)
        elif tipo == 'h2':
            doc.add_heading(contenido, level=1)
        elif tipo == 'h3':
            doc.add_heading(contenido, level=2)
        elif tipo == 'p':
            limpio, es_cursiva = cursiva_entera(contenido)
            p = doc.add_paragraph()
            if es_cursiva:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                r = p.add_run(limpio)
                r.italic = True
                r.font.color.rgb = RGBColor(0x44, 0x44, 0x44)
            else:
                escribir_runs(p, contenido)
        elif tipo == 'ul':
            for it in contenido:
                escribir_runs(doc.add_paragraph(style='List Bullet'), it)
        elif tipo == 'ol':
            for it in contenido:
                escribir_runs(doc.add_paragraph(style='List Number'), it)
        elif tipo == 'code':
            r = doc.add_paragraph().add_run(contenido)
            r.font.name = 'Courier New'
            r.font.size = Pt(10)
        elif tipo == 'table':
            n_cols = len(contenido[0])
            tabla = doc.add_table(rows=0, cols=n_cols)
            tabla.style = 'Light Grid Accent 1'
            for idx_fila, fila in enumerate(contenido):
                row = tabla.add_row()
                for idx_col in range(n_cols):
                    valor = fila[idx_col] if idx_col < len(fila) else ''
                    cell = row.cells[idx_col]
                    cell.text = ''
                    escribir_runs(cell.paragraphs[0], valor,
                                  negrita_base=(idx_fila == 0))
        elif tipo == 'hr':
            doc.add_paragraph()
    doc.core_properties.title = PARAMS['titulo_pdf']
    doc.core_properties.author = PARAMS['autor_pdf']
    doc.core_properties.language = 'en-US'
    doc.save(salida_docx)
    return salida_docx


# ==========================================================================
# PDF (reportlab, Letter)
# ==========================================================================
def _cabecera_pie(canv, doc_):
    canv.saveState()
    ancho, alto = PARAMS['pagesize']
    canv.setFont('Helvetica', 8)
    canv.setFillColor(colors.HexColor('#888888'))
    canv.drawString(2 * cm, alto - 1.3 * cm, PARAMS['cabecera'])
    canv.drawRightString(ancho - 2 * cm, 1.2 * cm,
                         PARAMS['pie'].format(n=doc_.page))
    canv.setStrokeColor(colors.HexColor('#DDDDDD'))
    canv.line(2 * cm, alto - 1.5 * cm, ancho - 2 * cm, alto - 1.5 * cm)
    canv.restoreState()


def _md_a_html(texto):
    texto = texto.replace('&', '&amp;').replace('<', '&lt;').replace('>', '&gt;')
    return re.sub(r'\*\*([^*]+)\*\*', r'<b>\1</b>', texto)


def construir_pdf(bloques, salida_pdf, titulo_meta):
    styles = getSampleStyleSheet()
    estilo_h1 = ParagraphStyle('H1c', parent=styles['Title'], fontName='Helvetica-Bold',
                               fontSize=22, leading=26, textColor=colors.HexColor(DARK),
                               spaceAfter=4, alignment=1)
    estilo_meta = ParagraphStyle('Meta', parent=styles['Normal'], fontName='Helvetica-Oblique',
                                 fontSize=11, textColor=colors.HexColor('#666666'),
                                 alignment=1, spaceAfter=16)
    estilo_h2 = ParagraphStyle('H2c', parent=styles['Heading2'], fontName='Helvetica-Bold',
                               fontSize=15, leading=19, textColor=colors.HexColor(GOLD),
                               spaceBefore=14, spaceAfter=8)
    estilo_h3 = ParagraphStyle('H3c', parent=styles['Heading3'], fontName='Helvetica-Bold',
                               fontSize=12.5, leading=16, textColor=colors.HexColor(DARK),
                               spaceBefore=10, spaceAfter=6)
    estilo_p = ParagraphStyle('Pc', parent=styles['Normal'], fontName='Helvetica',
                              fontSize=10.2, leading=14.5, spaceAfter=8,
                              alignment=4)
    estilo_li = ParagraphStyle('Lic', parent=estilo_p, spaceAfter=3)
    estilo_firma = ParagraphStyle('Firma', parent=estilo_p,
                                  fontName='Helvetica-Oblique', alignment=1,
                                  textColor=colors.HexColor('#444444'),
                                  spaceBefore=10)
    estilo_code = ParagraphStyle('Codec', parent=styles['Code'], fontName='Courier',
                                 fontSize=9, leading=12, backColor=colors.HexColor('#F5F5F5'),
                                 borderPadding=6, spaceAfter=10)

    ancho_pag = PARAMS['pagesize'][0]
    flujo = []
    primero_h1 = True
    for tipo, contenido in bloques:
        if tipo == 'h1':
            if primero_h1:
                flujo.append(Paragraph(_md_a_html(contenido), estilo_h1))
                flujo.append(Paragraph(_md_a_html(titulo_meta), estilo_meta))
                primero_h1 = False
            else:
                flujo.append(Paragraph(_md_a_html(contenido), estilo_h2))
        elif tipo == 'h2':
            flujo.append(Paragraph(_md_a_html(contenido), estilo_h2))
        elif tipo == 'h3':
            flujo.append(Paragraph(_md_a_html(contenido), estilo_h3))
        elif tipo == 'p':
            limpio, es_cursiva = cursiva_entera(contenido)
            flujo.append(Paragraph(_md_a_html(limpio),
                                   estilo_firma if es_cursiva else estilo_p))
        elif tipo == 'ul':
            items = [ListItem(Paragraph(_md_a_html(it), estilo_li), leftIndent=6)
                     for it in contenido]
            flujo.append(ListFlowable(items, bulletType='bullet', start='•',
                                      leftIndent=16, spaceAfter=8))
        elif tipo == 'ol':
            items = [ListItem(Paragraph(_md_a_html(it), estilo_li), leftIndent=6)
                     for it in contenido]
            flujo.append(ListFlowable(items, bulletType='1', leftIndent=16,
                                      spaceAfter=8))
        elif tipo == 'code':
            flujo.append(Preformatted(contenido, estilo_code))
        elif tipo == 'table':
            filas = contenido
            data = [[Paragraph(_md_a_html(c), ParagraphStyle(
                'Cellc', parent=estilo_p, fontSize=8.6, leading=11,
                alignment=0, spaceAfter=0,
                fontName='Helvetica-Bold' if fi == 0 else 'Helvetica'))
                for c in fila] for fi, fila in enumerate(filas)]
            n_cols = len(filas[0])
            col_w = (ancho_pag - 4 * cm) / n_cols
            t = Table(data, colWidths=[col_w] * n_cols, repeatRows=1)
            t.setStyle(TableStyle([
                ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor(GOLD)),
                ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
                ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#DDDDDD')),
                ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
                ('ROWBACKGROUNDS', (0, 1), (-1, -1),
                 [colors.white, colors.HexColor('#FAFAFA')]),
                ('LEFTPADDING', (0, 0), (-1, -1), 5),
                ('RIGHTPADDING', (0, 0), (-1, -1), 5),
                ('TOPPADDING', (0, 0), (-1, -1), 4),
                ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
            ]))
            flujo.append(t)
            flujo.append(Spacer(1, 10))
        elif tipo == 'hr':
            flujo.append(Spacer(1, 6))

    doc = SimpleDocTemplate(salida_pdf, pagesize=PARAMS['pagesize'],
                            topMargin=2.2 * cm, bottomMargin=2 * cm,
                            leftMargin=2 * cm, rightMargin=2 * cm,
                            title=PARAMS['titulo_pdf'],
                            author=PARAMS['autor_pdf'],
                            subject=PARAMS['asunto_pdf'],
                            creator=PARAMS['creador_pdf'],
                            keywords=PARAMS['keywords_pdf'])
    doc.build(flujo, onFirstPage=_cabecera_pie, onLaterPages=_cabecera_pie)
    return salida_pdf


# ==========================================================================
# Orquestación y gates
# ==========================================================================
def _abortar(msg):
    print('ABORTADO: ' + msg, file=sys.stderr)
    raise SystemExit(3)


def preparar(md_text):
    """Parsea, separa la línea de metadatos y pasa los gates 1 y 2 sobre el
    texto FUENTE. Devuelve (bloques_saneados, titulo_meta, informe)."""
    bloques = parsear(md_text)
    titulo_meta = PARAMS['meta_por_defecto']
    for tipo, contenido in bloques:
        if tipo == 'p' and PARAMS['marcador_meta'] in contenido:
            titulo_meta = re.sub(r'\*\*', '', contenido)
            break
    cuerpo = [b for b in bloques
              if not (b[0] == 'p' and re.sub(r'\*\*', '', b[1]) == titulo_meta)]

    fijos = [PARAMS['cabecera'], PARAMS['pie'], PARAMS['titulo_pdf'],
             PARAMS['asunto_pdf'], PARAMS['creador_pdf'], PARAMS['keywords_pdf'],
             titulo_meta]
    textos = list(_trozos(cuerpo)) + fijos

    # Gate 2 sobre la fuente
    restos = restos_espanol(textos)
    if restos:
        _abortar('restos de español en la fuente → ' +
                 ' | '.join('%s: …%s…' % r for r in restos[:25]) +
                 (' (+%d)' % (len(restos) - 25) if len(restos) > 25 else ''))

    # Gate 1 sobre la fuente
    quedan = restos_no_winansi(textos)
    if quedan:
        _abortar('caracteres que Helvetica no puede pintar → ' +
                 ', '.join('%r (U+%04X) x%d' % (c, ord(c), n)
                           for c, n in sorted(quedan.items())) +
                 '. Añádelos a MAPA_WINANSI o quítalos del .md.')
    n_mapeados = fuera_de_cp1252(textos)
    cuerpo = sanear_bloques(cuerpo)
    titulo_meta = sanear(titulo_meta)
    # Doble comprobación tras el saneado (no puede quedar nada)
    if restos_no_winansi(list(_trozos(cuerpo)) + [titulo_meta]) or \
            fuera_de_cp1252(list(_trozos(cuerpo)) + [titulo_meta]):
        _abortar('quedan caracteres fuera de WinAnsi tras el saneado')
    informe = {
        'bloques': len(bloques),
        'h2': sum(1 for t, _ in bloques if t == 'h2'),
        'tablas': sum(1 for t, _ in bloques if t == 'table'),
        'saneados': n_mapeados,
        'palabras': len(re.findall(r"[A-Za-zÀ-ÿ0-9$%'.,/-]+",
                                   re.sub(r'[#*|`>-]{2,}', ' ', md_text))),
    }
    return cuerpo, titulo_meta, informe


def verificar_pdf(ruta_pdf):
    """Gate 3: relee el PDF, cuenta páginas, pie en cada página y restos de
    español / glifos en lo que el cliente ve."""
    try:
        from pypdf import PdfReader
    except ImportError:                                   # pragma: no cover
        from PyPDF2 import PdfReader
    r = PdfReader(ruta_pdf)
    paginas = len(r.pages)
    sin_pie = []
    textos = []
    for i, p in enumerate(r.pages, 1):
        t = p.extract_text() or ''
        textos.append(t)
        if PARAMS['pie'].format(n=i) not in t:
            sin_pie.append(i)
    if sin_pie:
        _abortar('páginas sin pie «%s»: %s' % (PARAMS['pie'], sin_pie[:10]))
    todo = '\n'.join(textos)
    if '■' in todo:
        _abortar('el PDF contiene ■ (glifo no pintable)')
    restos = restos_espanol([todo])
    if restos:
        _abortar('restos de español en el PDF → ' +
                 ' | '.join('%s: …%s…' % x for x in restos[:15]))
    meta = r.metadata or {}
    if (meta.get('/Title') or '') != PARAMS['titulo_pdf']:
        _abortar('título del PDF inesperado: %r' % meta.get('/Title'))
    mb = r.pages[0].mediabox
    ancho, alto = float(mb.width), float(mb.height)
    if (round(ancho), round(alto)) != (612, 792):
        _abortar('el papel no es US Letter: %sx%s pt' % (ancho, alto))
    return paginas


def autotest():
    """Prueba en negativo: cada defecto inyectado tiene que hacer saltar su gate."""
    base = ('# How to Reduce Food Cost in 30 Days\n\n'
            '**A bonus guide from Food Cost Kit Pro by AI Chef Pro**\n\n'
            'Food cost is what you use, not what you buy. Visit '
            'https://aichef.pro/en/digital-products or johnguerrero.es. '
            'Café, purée and açaí are fine.\n')
    casos = {
        'signo €': 'It costs 12 €.',
        'comillas «»': 'The «Recipe Cost Card» tab.',
        'apertura ¿': '¿Te suena?',
        'tilde ES': 'Guía rápida.',
        'eñe': 'Diseño limpio.',
        'funcional de': 'Food cost de verdad.',
        'funcional que': 'The one que matters.',
        'oficio IVA': 'Prices without IVA.',
        'oficio escandallo': 'Open the escandallo.',
        'oficio merma': 'Measure the merma.',
        'prohibida entrée': 'A classic entrée.',
        'glifo ≈': 'About ≈ 30%.',
        'glifo CJK': 'Mix 鬼笔 well.',
        'glifo ⅓': 'Use ⅓ of a pint.',
    }
    try:
        preparar(base)
    except SystemExit:
        print('AUTOTEST: la base limpia no pasa (falso positivo)')
        return 1
    fallos = []
    for nombre, frase in casos.items():
        try:
            preparar(base + '\n' + frase + '\n')
            fallos.append(nombre)
        except SystemExit:
            pass
    detectados = len(casos) - len(fallos)
    print('AUTOTEST: %d/%d defectos detectados; base limpia sin falsos positivos'
          % (detectados, len(casos)))
    if fallos:
        print('NO detectados: ' + ', '.join(fallos))
        return 1
    return 0


def main():
    ap = argparse.ArgumentParser(description=__doc__.split('\n\n')[0])
    ap.add_argument('entrada', nargs='?')
    ap.add_argument('salida_pdf', nargs='?', help='ruta del PDF sin extensión')
    ap.add_argument('--docx', help='ruta del .docx de trabajo sin extensión')
    ap.add_argument('--solo-gates', action='store_true')
    ap.add_argument('--autotest', action='store_true')
    a = ap.parse_args()

    if a.autotest:
        raise SystemExit(autotest())
    if not a.entrada or (not a.salida_pdf and not a.solo_gates):
        ap.error('faltan argumentos')

    with open(a.entrada, encoding='utf-8') as f:
        md_text = f.read()
    cuerpo, titulo_meta, inf = preparar(md_text)
    print('GATES FUENTE OK: glifos 0 fuera de WinAnsi (%d saneados por mapa), '
          'restos de español 0' % inf['saneados'])
    if a.solo_gates:
        print('bloques=%(bloques)d h2=%(h2)d tablas=%(tablas)d palabras~%(palabras)d' % inf)
        return

    pdf_path = construir_pdf(cuerpo, a.salida_pdf + '.pdf', titulo_meta)
    paginas = verificar_pdf(pdf_path)
    print('OK: %s' % pdf_path)
    if a.docx:
        print('OK: %s' % construir_docx(cuerpo, a.docx + '.docx', titulo_meta))
    print('bloques=%(bloques)d h2=%(h2)d tablas=%(tablas)d palabras~%(palabras)d '
          'saneados=%(saneados)d' % inf + ' paginas=%d papel=Letter pie=OK' % paginas)


if __name__ == '__main__':
    main()
