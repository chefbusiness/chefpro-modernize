#!/usr/bin/env python3
"""
calculadora_sus_es.py — Plan de Negocio: Sushi Bar ES 1.0: `calculadora-pescado-arroz-escandallo-sushi-bar.xlsx`.
Hojas `Pescado` · `Arroz` · `Escandallo` · `Instrucciones`. Gemelo de código: `plan-polleria/calculadora_pol_es.py`
(helpers, inject_cache, autotest con pycel y un `calcular()` en Python, protección sin contraseña, verdes desbloqueadas, A4,
docProps AI Chef Pro). Todos los valores por defecto son «estimación del kit, editable» (F1-research-sushi.md §2 y §8).
  · Pescado: 3 pescados + 1 libre (salmón entero de acuicultura, lomo comprado limpio, atún en saku) → coste por kg útil =
    precio ÷ (rendimiento + recorte) ÷ (1 − merma de vitrina), coste por loncha, corte y roll, mano de obra de limpieza por kg útil
    y «el entero te compensa mientras limpiarlo cueste menos de X € por kg útil». Desplegable informativo de congelación.
  · Arroz: kg de shari por kg de arroz crudo, coste por kg de shari, nigiris por kilo de arroz crudo, coste del shari y del nori
    por pieza y kg de arroz crudo al día.
  · Escandallo: 14 productos × (pescado, neta, shari, nori, otros, PVP, unidades a la semana) → coste, food cost con semáforo,
    PVP sugerido, margen por canal (sala · para llevar · plataforma), puesto por margen y food cost ponderado.
Valores que otros scripts leen de la CACHÉ del xlsx (dict `EXPUESTOS`, función `valores_expuestos(ruta)`).
Corre en el VPS (openpyxl + pycel):

    /root/venv-guias/bin/python calculadora_sus_es.py --salida <dir>      # genera + inject_cache + autotest
    /root/venv-guias/bin/python calculadora_sus_es.py --salida <dir> --autotest

Sesión Claude Code, 31-oct-2026.
"""
import argparse
import math
import os
import re
import sys
from collections import OrderedDict

import openpyxl
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill, Protection
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

AQUI = os.path.dirname(os.path.abspath(__file__))
PD = os.path.dirname(AQUI)
sys.path.insert(0, AQUI)
import sus_version                                                                 # noqa: E402,F401  (parchea motor: 1.0 · noviembre 2026)
import motor                                                                       # noqa: E402

PID = 'plan-negocio-sushi-bar'
FICHERO = 'calculadora-pescado-arroz-escandallo-sushi-bar.xlsx'
TITULO = 'Calculadora del pescado, el arroz y el escandallo'
HP, HA, HE, HI = 'Pescado', 'Arroz', 'Escandallo', 'Instrucciones'
VERDE = PatternFill('solid', fgColor='E8F5E9')
GRIS = PatternFill('solid', fgColor='F2F2F2')
ROJO = PatternFill('solid', bgColor='F8CBAD', fgColor='F8CBAD')
AMBAR = PatternFill('solid', bgColor='FFE699', fgColor='FFE699')
OK = PatternFill('solid', bgColor='C6E0B4', fgColor='C6E0B4')
FUENTE = 'Calibri'
EUR, ENT, PCT, DEC, GR, EUR3 = '#,##0.00 €', '#,##0', '0.0%', '#,##0.0', '#,##0.0', '#,##0.000 €'
FC_OK, FC_VIG, FC_ALTO = 'En objetivo', 'Vigila el coste', 'Por encima del objetivo'
SEM_ENT, SEM_LOMO = 'El entero compensa', 'Compra el lomo limpio'
EST = ' Estimación del kit, editable.'

# ---------------------------------------------------------------------------------------------------------------------
# Datos (F1-research §2 y §8). Importes sin IVA salvo el PVP de carta.
# ---------------------------------------------------------------------------------------------------------------------
NPE = 4
PC = [get_column_letter(2 + i) for i in range(NPE)]                                # B..E
#: hoja Pescado, bloque A: (nombre, precio €/kg, rendimiento del lomo, recorte reaprovechado)
PESCADOS = [
    ('Salmón entero de acuicultura', 9.50, 0.55, 0.15),
    ('Lomo de salmón comprado limpio', 21.80, 0.92, 0.03),
    ('Atún aleta amarilla en saku ultracongelado', 26.00, 0.90, 0.05),
    ('Otro pescado (escribe el tuyo)', 18.00, 0.80, 0.05),
]
P_ROW = OrderedDict([('pr', 6), ('rd', 7), ('rc', 8)])
P_NOTAS = OrderedDict([
    ('pr', 'Salmón noruego entero de acuicultura: unos 11,95 €/kg en minorista y 6-7,5 € en origen; lomo o filete congelado, de 20,7 a 21,8 €/kg '
           '(Makro); el atún aleta amarilla en saku ultracongelado no tiene precio publicado en España. El kit pone 9,50 · 21,80 · 26,00 €/kg.' + EST),
    ('rd', 'Parte del pescado entero que queda como lomo limpio para sushi: de 50 a 60 % en el salmón entero. Mídelo con 3 salmones '
           '(hoja «Instrucciones»). En el atún en saku es la parte que se aprovecha al cortar.' + EST),
    ('rc', 'Recorte que reaprovechas (tartar, maki, poke), en % del peso comprado. Lo que tiras no cuenta aquí.' + EST),
])
OPC = [  # desplegable informativo: (opción, recordatorio)
    ('Congelo yo: −20 °C ≥ 24 h o −35 °C ≥ 15 h',
     'RD 1021/2022, art. 8: −20 °C o menos durante al menos 24 h, o −35 °C o menos durante al menos 15 h, en la totalidad del producto. '
     'Registra temperatura y horas de cada lote.'),
    ('Llega congelado con documento del proveedor',
     'Guarda el documento del proveedor de cada lote: tiene que acreditar que el pescado se ha congelado a esas temperaturas y horas.'),
    ('Acuicultura con documento escrito',
     'La acuicultura queda exenta solo con el documento escrito del proveedor. Archívalo con el albarán y pídelo de nuevo si cambia de origen.'),
]
#: bloque B (entradas comunes de la hoja Pescado): (clave, rótulo, formato, valor, constante, nota)
ENT_PE = [
    ('m', 'Merma de vitrina (% del pescado cortado que no se vende)', PCT, 0.05, False,
     'Lo que se queda en la vitrina o en el pase y se tira al cierre.' + EST),
    ('gn', 'Gramos de neta por loncha de nigiri (g)', DEC, 15.0, False, 'Criterio del plan: pesa 10 lonchas y usa la media.' + EST),
    ('gs', 'Gramos por corte de sashimi (g)', DEC, 15.0, False, 'Pesa 10 cortes y usa la media.' + EST),
    ('gr', 'Gramos de pescado por roll (g)', DEC, 40.0, False, 'Un uramaki o un maki de 8 piezas con pescado dentro.' + EST),
    ('tl', 'Minutos de limpieza de un pescado entero (min)', DEC, 25.0, False,
     'Sin fuente: mídelo. Cronometra la limpieza de 3 piezas y usa la media.' + EST),
    ('hh', 'Coste por hora del itamae (€/h)', EUR, 22.0, False,
     '2.200 € brutos × 14 pagas + 33 % de cargas ÷ 1.840 h ≈ 22 €/h.' + EST),
    ('pe', 'Peso de una pieza entera (kg)', '#,##0.00', 5.0, False, 'Salmón entero de acuicultura de unos 5 kg.' + EST),
    ('cong', '¿Congelación anti-anisakis o documento? (desplegable)', 'General', OPC[0][0], False,
     'Informativo: elige tu caso y mira el recordatorio de abajo. No cambia ningún cálculo.'),
    ('KG', 'Gramos por kilo (constante)', ENT, 1000, True, 'Constante: un kilo son 1.000 gramos. No se edita.'),
    ('MH', 'Minutos por hora (constante)', ENT, 60, True, 'Constante: una hora son 60 minutos. No se edita.'),
]
PE_M0 = 11
M_PE = OrderedDict((k[0], PE_M0 + i) for i, k in enumerate(ENT_PE))                # 11..20
PE_RC = PE_M0 + len(ENT_PE) + 1                                                    # 22: cabecera del bloque C
SAL_C = [
    ('tot', 'Parte aprovechada (rendimiento + recorte)', PCT, '{rd}+{rc}', 'Lomo + recorte que reaprovechas.'),
    ('cku', 'COSTE POR KG ÚTIL (€)', EUR, '{pr}/{tot}/(1-{m})', 'Precio ÷ (rendimiento + recorte) ÷ (1 − merma de vitrina).'),
    ('c_nig', 'Coste de una loncha de nigiri (€)', EUR3, '{cku}*{gn}/{KG}', 'Coste por kg útil × gramos ÷ 1.000.'),
    ('c_sas', 'Coste de un corte de sashimi (€)', EUR3, '{cku}*{gs}/{KG}', 'Coste por kg útil × gramos ÷ 1.000.'),
    ('c_roll', 'Coste del pescado de un roll (€)', EUR3, '{cku}*{gr}/{KG}', 'Coste por kg útil × gramos ÷ 1.000.'),
]
S_C = OrderedDict((k[0], PE_RC + 1 + i) for i, k in enumerate(SAL_C))             # 23..27
PE_RD = PE_RC + len(SAL_C) + 2                                                     # 29: cabecera del bloque D
SAL_D = [
    ('lab_pz', 'Coste de limpiar una pieza entera (€)', EUR, '{tl}/{MH}*{hh}', 'Minutos de limpieza ÷ 60 × coste por hora del itamae.'),
    ('kg_pz', 'Kilos útiles de una pieza entera (kg)', '#,##0.00', '{pe}*{rd0}', 'Peso de la pieza × rendimiento del salmón entero.'),
    ('lab_kg', 'MANO DE OBRA DE LIMPIEZA POR KG ÚTIL (€)', EUR, '{lab_pz}/{kg_pz}', 'Coste de limpiar una pieza ÷ kilos útiles.'),
    ('c_ent', 'Coste por kg útil del entero con la limpieza incluida (€)', EUR, '{cku0}+{lab_kg}', 'Coste por kg útil + mano de obra por kg útil.'),
    ('X', 'EL ENTERO TE COMPENSA MIENTRAS LIMPIARLO CUESTE MENOS DE (€ por kg útil)', EUR, '{cku1}-{cku0}',
     'Coste por kg útil del lomo comprado − coste por kg útil del entero.'),
    ('sem', 'Veredicto con tu mano de obra', 'General', 'IF({lab_kg}<{X},"%s","%s")' % (SEM_ENT, SEM_LOMO),
     'El entero compensa si limpiarlo cuesta menos que esa diferencia.'),
    ('rem', 'Recordatorio de tu caso de congelación', 'General', 'INDEX({rem_r},MATCH({cong},{opc_r},0))', 'Texto de la lista de abajo.'),
]
S_D = OrderedDict((k[0], PE_RD + 1 + i) for i, k in enumerate(SAL_D))             # 30..36
PE_RO = PE_RD + len(SAL_D) + 2                                                     # 38: cabecera de la lista
PE_O0 = PE_RO + 1                                                                  # 39..41
PE_O1 = PE_O0 + len(OPC) - 1

#: hoja Arroz: (clave, rótulo, formato, valor, constante, nota)
ENT_A = [
    ('ak', 'Arroz crudo para sushi (€/kg, sin IVA)', EUR, 1.85, False, 'Entre 1,81 €/kg (METRO Chef, 5 kg) y 1,85 €/kg (DACSA, 5 kg); otra marca, a 2,50 €/kg en saco de 20 kg.' + EST),
    ('f', 'Factor de arroz crudo a cocido (×)', '#,##0.00', 2.2, False, 'Un kilo de arroz crudo da 2,2 kg de arroz cocido.' + EST),
    ('z', 'Sushi-zu (aliño de vinagre) por kg de arroz cocido (ml)', ENT, 100, False, 'Vinagre de arroz, azúcar y sal: unos 100 ml por kilo de arroz cocido.' + EST),
    ('zp', 'Sushi-zu (€/l, sin IVA)', EUR, 4.40, False, 'Vinagre, azúcar y sal.' + EST),
    ('g_nig', 'Gramos de shari por nigiri (g)', DEC, 20.0, False, 'Shari es el arroz ya aliñado. Pesa 10 bolas y usa la media.' + EST),
    ('g_hos', 'Gramos de shari por hosomaki (g)', DEC, 70.0, False, 'Un rollo fino de 6 piezas.' + EST),
    ('g_ura', 'Gramos de shari por uramaki (g)', DEC, 120.0, False, 'Un rollo de 8 piezas.' + EST),
    ('g_tem', 'Gramos de shari por temaki (g)', DEC, 60.0, False, 'Un cono.' + EST),
    ('g_pok', 'Gramos de shari por poke (g)', DEC, 150.0, False, 'Un bol.' + EST),
    ('nori', 'Nori (€ por hoja, sin IVA)', EUR3, 0.26, False, 'Nori Gold, 100 hojas por 25,95 € = 0,26 €/hoja.' + EST),
    ('n_nig', 'Hojas de nori por nigiri', '#,##0.00', 0.05, False, 'La tira que ata el nigiri.' + EST),
    ('n_hos', 'Hojas de nori por hosomaki', '#,##0.00', 0.5, False, 'Media hoja por rollo.' + EST),
    ('n_ura', 'Hojas de nori por uramaki', '#,##0.00', 0.5, False, 'Media hoja por rollo.' + EST),
    ('n_tem', 'Hojas de nori por temaki', '#,##0.00', 0.5, False, 'Media hoja por cono.' + EST),
    ('n_pok', 'Hojas de nori por poke', '#,##0.00', 0.1, False, 'Un toque de nori cortado por encima.' + EST),
    ('mt', 'Shari que se tira al cierre (%)', PCT, 0.05, False, 'Lo que sobra y no se puede guardar.' + EST),
    ('cd', 'Comensales al día', ENT, 79, False, 'El del caso del plan de negocio. Cámbialo por el tuyo.' + EST),
    ('gc', 'Gramos de shari por comensal (g)', DEC, 150.0, False, 'Piezas por comensal × gramos por pieza.' + EST),
    ('KG', 'Gramos por kilo (constante)', ENT, 1000, True, 'Constante: un kilo son 1.000 gramos. No se edita.'),
    ('ML', 'Mililitros por litro (constante)', ENT, 1000, True, 'Constante: un litro son 1.000 ml. No se edita.'),
]
SAL_A = [
    ('kgs', 'KG DE SHARI POR KG DE ARROZ CRUDO (kg)', '#,##0.00', '{f}*(1+{z}/{ML})', 'Factor crudo a cocido × (1 + sushi-zu en litros por kg cocido).'),
    ('cs', 'COSTE POR KG DE SHARI (€)', EUR, '({ak}+{f}*{z}/{ML}*{zp})/{kgs}', '(Arroz + sushi-zu de un kilo de arroz crudo) ÷ kg de shari.'),
    ('csm', 'Coste por kg de shari con la merma del cierre (€)', EUR, '{cs}/(1-{mt})', 'Coste por kg de shari ÷ (1 − shari que se tira). Alimenta el escandallo.'),
    ('nigkg', 'NIGIRIS POR KILO DE ARROZ CRUDO (piezas)', ENT, '{KG}*{kgs}/{g_nig}', '1.000 × kg de shari por kg crudo ÷ gramos por nigiri.'),
    ('c_nig', 'Coste del shari de un nigiri (€)', EUR3, '{g_nig}*{csm}/{KG}', 'Gramos × coste por kg de shari con merma ÷ 1.000.'),
    ('c_hos', 'Coste del shari de un hosomaki (€)', EUR3, '{g_hos}*{csm}/{KG}', 'Gramos × coste por kg de shari con merma ÷ 1.000.'),
    ('c_ura', 'Coste del shari de un uramaki (€)', EUR3, '{g_ura}*{csm}/{KG}', 'Gramos × coste por kg de shari con merma ÷ 1.000.'),
    ('c_tem', 'Coste del shari de un temaki (€)', EUR3, '{g_tem}*{csm}/{KG}', 'Gramos × coste por kg de shari con merma ÷ 1.000.'),
    ('c_pok', 'Coste del shari de un poke (€)', EUR3, '{g_pok}*{csm}/{KG}', 'Gramos × coste por kg de shari con merma ÷ 1.000.'),
    ('o_nig', 'Coste del nori de un nigiri (€)', EUR3, '{n_nig}*{nori}', 'Hojas × precio de la hoja.'),
    ('o_hos', 'Coste del nori de un hosomaki (€)', EUR3, '{n_hos}*{nori}', 'Hojas × precio de la hoja.'),
    ('o_ura', 'Coste del nori de un uramaki (€)', EUR3, '{n_ura}*{nori}', 'Hojas × precio de la hoja.'),
    ('o_tem', 'Coste del nori de un temaki (€)', EUR3, '{n_tem}*{nori}', 'Hojas × precio de la hoja.'),
    ('o_pok', 'Coste del nori de un poke (€)', EUR3, '{n_pok}*{nori}', 'Hojas × precio de la hoja.'),
    ('shari_dia', 'Shari que vendes al día (kg)', '#,##0.00', '{cd}*{gc}/{KG}', 'Comensales × gramos por comensal ÷ 1.000.'),
    ('arroz_dia', 'KG DE ARROZ CRUDO AL DÍA', '#,##0.00', '{shari_dia}/{kgs}', 'Shari que vendes ÷ kg de shari por kg de arroz crudo.'),
    ('arroz_dia_m', 'Kg de arroz crudo al día contando lo que se tira', '#,##0.00', '{arroz_dia}/(1-{mt})', 'Arroz crudo al día ÷ (1 − shari que se tira).'),
]
M_A, M_RA, S_A = None, None, None


def _filas(ent, sal, h0=5):
    m = OrderedDict((k[0], h0 + 1 + i) for i, k in enumerate(ent))
    ra = h0 + 1 + len(ent) + 1
    s = OrderedDict((k[0], ra + 1 + i) for i, k in enumerate(sal))
    return m, ra, s


M_A, M_RA, S_A = _filas(ENT_A, SAL_A)

#: hoja Escandallo: datos generales
ESC_IN = [
    ('iva', 'IVA de la comida (%)', PCT, 0.10, 'Comida preparada, en sala, para llevar y a domicilio: 10 % (Ley 37/1992, art. 91).'),
    ('obj', 'Food cost objetivo (% sobre el precio sin IVA)', PCT, 0.30, 'Objetivo de la carta del sushi bar.' + EST),
    ('tol', 'Margen de vigilancia sobre el objetivo (puntos)', PCT, 0.05, 'Hasta este margen por encima del objetivo, el semáforo dice «Vigila el coste».'),
    ('cl', 'Caja para llevar por pedido (€, sin IVA)', EUR, 0.40, 'Bandeja, cierre, bolsa, salsas y palillos de un pedido para llevar.' + EST),
    ('cp', 'Caja de plataforma por pedido (€, sin IVA)', EUR, 0.90, 'Envase de reparto más sellado.' + EST),
    ('rec', 'Recargo de precio en la plataforma (%)', PCT, 0.15, 'Lo que subes el precio en la plataforma para compensar.' + EST),
    ('com', 'Comisión de la plataforma de reparto (%)', PCT, 0.30, 'Sobre el precio sin IVA del pedido. Revisa tu contrato.' + EST),
]
E = OrderedDict((k[0], 5 + i) for i, k in enumerate(ESC_IN))
E_SH, E_NO = 5 + len(ESC_IN), 5 + len(ESC_IN) + 1                                # enlaces a la hoja Arroz
E_CAB = E_NO + 2
#: producto: (nombre, nº de pescado 0-4, g de neta, g de shari, hojas de nori, otros ingredientes €, PVP con IVA de sala, unidades a la semana)
PRODS = [
    ('Nigiri de salmón (2 piezas)', 1, 30, 40, 0.10, 0.03, 5.20, 110),
    ('Nigiri de atún (2 piezas)', 3, 30, 40, 0.10, 0.03, 6.00, 70),
    ('Sashimi de salmón (5 cortes)', 1, 75, 0, 0.0, 0.04, 7.50, 50),
    ('Hosomaki de salmón (6 piezas)', 1, 25, 70, 0.50, 0.05, 5.50, 20),
    ('Uramaki California (8 piezas)', 0, 0, 120, 0.50, 0.52, 9.60, 35),
    ('Uramaki de salmón y aguacate (8 piezas)', 1, 40, 120, 0.50, 0.45, 11.50, 35),
    ('Uramaki de langostino en tempura (8 piezas)', 0, 0, 120, 0.50, 1.37, 11.50, 40),
    ('Temaki de salmón', 1, 40, 60, 0.50, 0.40, 6.90, 30),
    ('Poke de salmón', 1, 100, 150, 0.10, 0.76, 12.90, 45),
    ('Combinado de 12 piezas', 1, 100, 190, 0.50, 0.45, 14.90, 25),
    ('Combinado de 24 piezas', 1, 220, 380, 1.00, 0.33, 29.00, 15),
    ('Gyozas (6 piezas)', 0, 0, 0, 0.0, 1.10, 7.00, 60),
    ('Edamame (ración)', 0, 0, 0, 0.0, 0.70, 4.50, 60),
    ('Mochi (2 piezas)', 0, 0, 0, 0.0, 1.25, 5.50, 45),
]
NPR = len(PRODS)
J_NIG, J_C24 = 0, 10
E_P0 = E_CAB + 1
E_P1 = E_P0 + NPR - 1
E_T0 = E_P1 + 2
TOT = [
    ('ing', 'Ventas de la semana sin IVA (€)', EUR, 'SUMPRODUCT(%s,%s)'),
    ('cos', 'Coste de la comida vendida en la semana (€)', EUR, 'SUMPRODUCT(%s,%s)'),
    ('fcp', 'FOOD COST PONDERADO (%)', PCT, None),
    ('semp', 'Semáforo del food cost ponderado', 'General', None),
]
T = OrderedDict((k[0], E_T0 + 1 + i) for i, k in enumerate(TOT))
E_FIN = E_T0 + len(TOT) + 2
# columnas de la tabla
CN, CPF, CNE, CSH, CNO, COT, CPVP, CUD = 'A', 'B', 'C', 'D', 'E', 'F', 'G', 'H'
CCO, CPS, CFC, CSE, CSG, CMS, CML, CMP, CRK = 'I', 'J', 'K', 'L', 'M', 'N', 'O', 'P', 'Q'
SALP = [('c_tot', CCO, EUR3), ('pvp_sin', CPS, EUR), ('fc', CFC, PCT), ('sem', CSE, 'General'), ('pvp_sug', CSG, EUR),
        ('sala', CMS, EUR), ('llevar', CML, EUR), ('plat', CMP, EUR), ('rank', CRK, ENT)]

EXPUESTOS = OrderedDict([
    ('COSTE_KG_UTIL_SALMON', (HP, 'B%d' % S_C['cku'])),
    ('COSTE_KG_UTIL_LOMO', (HP, 'C%d' % S_C['cku'])),
    ('COSTE_KG_UTIL_ATUN', (HP, 'D%d' % S_C['cku'])),
    ('KG_SHARI_KG_ARROZ', (HA, 'B%d' % S_A['kgs'])),
    ('NIGIRIS_KG_ARROZ', (HA, 'B%d' % S_A['nigkg'])),
    ('ARROZ_DIA_KG', (HA, 'B%d' % S_A['arroz_dia'])),
    ('FC_NIGIRI_SALMON', (HE, '%s%d' % (CFC, E_P0 + J_NIG))),
    ('FC_COMBINADO_24', (HE, '%s%d' % (CFC, E_P0 + J_C24))),
    ('MARGEN_SALA_COMB24', (HE, '%s%d' % (CMS, E_P0 + J_C24))),
    ('MARGEN_PLATAFORMA_COMB24', (HE, '%s%d' % (CMP, E_P0 + J_C24))),
    ('FC_PONDERADO', (HE, 'B%d' % T['fcp'])),
])


def valores_defecto():
    """Todas las entradas por defecto en un solo dict plano (claves = las de las celdas verdes)."""
    v = OrderedDict()
    for i, (_n, pr, rd, rc) in enumerate(PESCADOS):
        v['pr%d' % i], v['rd%d' % i], v['rc%d' % i] = pr, rd, rc
    for ent in (ENT_PE, ENT_A):
        for k, _r, _f2, val, _c, _n in ent:
            v[k] = val
    for k, _r, _f2, val, _n in ESC_IN:
        v[k] = val
    for j, (_nom, pf, ne, sh, no, ot, pvp, ud) in enumerate(PRODS):
        v['pf%d' % j], v['ne%d' % j], v['sh%d' % j], v['no%d' % j] = pf, ne, sh, no
        v['ot%d' % j], v['pvp%d' % j], v['ud%d' % j] = ot, pvp, ud
    return v


# ---------------------------------------------------------------------------------------------------------------------
# Cálculo en Python (espejo de las fórmulas) — fuente del autotest
# ---------------------------------------------------------------------------------------------------------------------
def calcular(v):
    r = OrderedDict()
    cols = []
    for i in range(NPE):
        c = OrderedDict()
        c['tot'] = v['rd%d' % i] + v['rc%d' % i]
        c['cku'] = v['pr%d' % i] / c['tot'] / (1 - v['m'])
        c['c_nig'] = c['cku'] * v['gn'] / v['KG']
        c['c_sas'] = c['cku'] * v['gs'] / v['KG']
        c['c_roll'] = c['cku'] * v['gr'] / v['KG']
        cols.append(c)
    r['Pe'] = cols
    d = OrderedDict()
    d['lab_pz'] = v['tl'] / v['MH'] * v['hh']
    d['kg_pz'] = v['pe'] * v['rd0']
    d['lab_kg'] = d['lab_pz'] / d['kg_pz']
    d['c_ent'] = cols[0]['cku'] + d['lab_kg']
    d['X'] = cols[1]['cku'] - cols[0]['cku']
    d['sem'] = SEM_ENT if d['lab_kg'] < d['X'] else SEM_LOMO
    d['rem'] = OPC[[o for o, _t in OPC].index(v['cong'])][1]
    r['D'] = d
    a = OrderedDict()
    a['kgs'] = v['f'] * (1 + v['z'] / v['ML'])
    a['cs'] = (v['ak'] + v['f'] * v['z'] / v['ML'] * v['zp']) / a['kgs']
    a['csm'] = a['cs'] / (1 - v['mt'])
    a['nigkg'] = v['KG'] * a['kgs'] / v['g_nig']
    for t in ('nig', 'hos', 'ura', 'tem', 'pok'):
        a['c_' + t] = v['g_' + t] * a['csm'] / v['KG']
    for t in ('nig', 'hos', 'ura', 'tem', 'pok'):
        a['o_' + t] = v['n_' + t] * v['nori']
    a['shari_dia'] = v['cd'] * v['gc'] / v['KG']
    a['arroz_dia'] = a['shari_dia'] / a['kgs']
    a['arroz_dia_m'] = a['arroz_dia'] / (1 - v['mt'])
    r['A'] = a
    r['shari'], r['nori'] = a['csm'], v['nori']
    kgc = [0.0] + [c['cku'] for c in cols]
    ps = []
    for j in range(NPR):
        q = OrderedDict()
        q['c_tot'] = v['ne%d' % j] / v['KG'] * kgc[int(v['pf%d' % j])] + v['sh%d' % j] / v['KG'] * r['shari'] + v['no%d' % j] * r['nori'] + v['ot%d' % j]
        q['pvp_sin'] = v['pvp%d' % j] / (1 + v['iva'])
        q['fc'] = q['c_tot'] / q['pvp_sin']
        q['sem'] = FC_OK if q['fc'] <= v['obj'] else (FC_VIG if q['fc'] <= v['obj'] + v['tol'] else FC_ALTO)
        q['pvp_sug'] = q['c_tot'] / v['obj'] * (1 + v['iva'])
        q['sala'] = q['pvp_sin'] - q['c_tot']
        q['llevar'] = q['sala'] - v['cl']
        q['plat'] = q['pvp_sin'] * (1 + v['rec']) * (1 - v['com']) - q['c_tot'] - v['cp']
        ps.append(q)
    for q in ps:
        q['rank'] = 1 + sum(1 for o in ps if o['sala'] > q['sala'])
    r['Pr'] = ps
    t = OrderedDict()
    t['ing'] = sum(q['pvp_sin'] * v['ud%d' % j] for j, q in enumerate(ps))
    t['cos'] = sum(q['c_tot'] * v['ud%d' % j] for j, q in enumerate(ps))
    t['fcp'] = 0 if t['ing'] == 0 else t['cos'] / t['ing']
    t['semp'] = FC_OK if t['fcp'] <= v['obj'] else (FC_VIG if t['fcp'] <= v['obj'] + v['tol'] else FC_ALTO)
    r['T'] = t
    return r


# ---------------------------------------------------------------------------------------------------------------------
# Helpers (de la Taquería y la Pollería)
# ---------------------------------------------------------------------------------------------------------------------
def _f(ws, coord, v, b=False, sz=11, i=False, fmt=None, wrap=False, fill=None, locked=True):
    c = ws[coord]
    c.value = v
    c.font = Font(name=FUENTE, size=sz, bold=b, italic=i)
    if fmt:
        c.number_format = fmt
    c.alignment = Alignment(wrap_text=wrap, vertical='top')
    if fill:
        c.fill = fill
    c.protection = Protection(locked=locked)
    return c


def _nota(ws, fila, texto, desde='E', hasta='F', chars=80.0):
    if not texto:
        return
    _f(ws, '%s%d' % (desde, fila), texto, wrap=True)
    if desde != hasta:
        ws.merge_cells('%s%d:%s%d' % (desde, fila, hasta, fila))
    ws.row_dimensions[fila].height = max(ws.row_dimensions[fila].height or 0, 16, 15 * math.ceil(len(texto) / chars))


def _imprimir(ws, horizontal=False):
    ws.page_setup.paperSize = 9                                                   # A4
    ws.page_setup.orientation = 'landscape' if horizontal else 'portrait'
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.oddFooter.center.text = 'aichef.pro/%s · página &P de &N' % PID
    ws.print_options.horizontalCentered = True
    ws.protection.sheet = True                                                    # sin contraseña
    ws.sheet_view.showGridLines = False


def _dv(ws, tipo, celdas, f1=None, f2=None, op=None, titulo='Valor no válido', msg=''):
    dv = DataValidation(type=tipo, formula1=f1, formula2=f2, operator=op, allow_blank=False, showErrorMessage=True,
                        errorTitle=titulo, error=msg)
    ws.add_data_validation(dv)
    for c in celdas:
        dv.add(c)


def a_excel(expr, ref):
    return re.sub(r'\{(\w+)\}', lambda m: ref[m.group(1)], expr)


def _bloque(ws, titulo_in, ent, m, ra, sal, s, neg, chars=105.0):
    """Escribe el bloque de entradas (verdes/grises) y el de resultados de una hoja de dos bloques."""
    _f(ws, 'A5', titulo_in, b=True, fill=GRIS)
    _f(ws, 'B5', 'Valor', b=True, fill=GRIS)
    _f(ws, 'C5', 'Notas', b=True, fill=GRIS)
    for k, rot, fmt, val, const, nota in ent:
        f = m[k]
        _f(ws, 'A%d' % f, rot)
        if const:
            _f(ws, 'B%d' % f, val, fmt=fmt, fill=GRIS)
        else:
            _f(ws, 'B%d' % f, val, fmt=fmt, fill=VERDE, locked=False)
        _nota(ws, f, nota, desde='C', hasta='I', chars=chars)
    ref = dict((k, '$B$%d' % f) for k, f in m.items())
    ref.update((k, '$B$%d' % f) for k, f in s.items())
    for col, t in zip('ABC', ('RESULTADOS', 'Valor', 'Notas')):
        _f(ws, '%s%d' % (col, ra), t, b=True, fill=GRIS)
    for k, rot, fmt, expr, nota in sal:
        f = s[k]
        _f(ws, 'A%d' % f, rot, b=k in neg)
        _f(ws, 'B%d' % f, '=' + a_excel(expr, ref), fmt=fmt, b=k in neg, wrap=(fmt == 'General'))
        _nota(ws, f, nota, desde='C', hasta='I', chars=chars)


def _semaforo(ws, celda, ok, ko):
    m = re.match(r'([A-Z]+)(\d+)', celda)
    ws.conditional_formatting.add(celda, FormulaRule(formula=['$%s$%s="%s"' % (m.group(1), m.group(2), ok)], fill=OK))
    ws.conditional_formatting.add(celda, FormulaRule(formula=['$%s$%s="%s"' % (m.group(1), m.group(2), ko)], fill=ROJO))


# ---------------------------------------------------------------------------------------------------------------------
# Hojas
# ---------------------------------------------------------------------------------------------------------------------
def hoja_pescado(wb):
    ws = wb.active
    ws.title = HP
    for col, w in (('A', 62), ('B', 19), ('C', 19), ('D', 19), ('E', 19), ('F', 72)):
        ws.column_dimensions[col].width = w
    _f(ws, 'A1', 'CALCULADORA DEL PESCADO — COSTE POR KILO ÚTIL, POR LONCHA, POR CORTE Y POR ROLL', b=True, sz=14)
    _f(ws, 'A2', 'Escribe solo en las celdas VERDES: todo lo demás se calcula. Importes sin IVA. Todos los valores por defecto son una '
                 'estimación del kit, editable: cámbialos por los de tu proveedor y, sobre todo, por lo que midas en tus propios pescados.', wrap=True)
    ws.merge_cells('A2:F2')
    ws.row_dimensions[2].height = 32
    _f(ws, 'A3', 'AVISO FIJO. Anisakis (RD 1021/2022, art. 8): el pescado que se sirve crudo o casi crudo se congela a −20 °C o menos durante '
                 'al menos 24 horas, o a −35 °C o menos durante al menos 15 horas, en la totalidad del producto; la acuicultura queda exenta '
                 'solo con documento escrito del proveedor, y hay que informar al cliente en la carta o en un cartel (art. 8.2). Los minutos de '
                 'limpieza (25) no tienen fuente: mídelos.', wrap=True, b=True, fill=GRIS)
    ws.merge_cells('A3:F3')
    ws.row_dimensions[3].height = 64
    # A. tres pescados y uno libre
    _f(ws, 'A5', 'A. TRES PESCADOS Y UNO LIBRE', b=True, fill=GRIS)
    for i, (nom, pr, rd, rc) in enumerate(PESCADOS):
        _f(ws, '%s5' % PC[i], nom, b=True, fill=VERDE, wrap=True, locked=False)
    _f(ws, 'F5', 'Notas', b=True, fill=GRIS)
    ws.row_dimensions[5].height = 62
    for k, rot, fmt, idx in (('pr', 'Precio de compra (€/kg, sin IVA)', EUR, 1), ('rd', 'Rendimiento: lomo limpio sobre el peso comprado', PCT, 2),
                             ('rc', 'Recorte que reaprovechas (% del peso comprado)', PCT, 3)):
        f = P_ROW[k]
        _f(ws, 'A%d' % f, rot)
        for i, p in enumerate(PESCADOS):
            _f(ws, '%s%d' % (PC[i], f), p[idx], fmt=fmt, fill=VERDE, locked=False)
        _nota(ws, f, P_NOTAS[k], desde='F', hasta='F', chars=75.0)
    # B. datos comunes
    _f(ws, 'A10', 'B. DATOS COMUNES', b=True, fill=GRIS)
    _f(ws, 'B10', 'Valor', b=True, fill=GRIS)
    _f(ws, 'C10', 'Notas', b=True, fill=GRIS)
    for k, rot, fmt, val, const, nota in ENT_PE:
        f = M_PE[k]
        _f(ws, 'A%d' % f, rot)
        if const:
            _f(ws, 'B%d' % f, val, fmt=fmt, fill=GRIS)
        else:
            _f(ws, 'B%d' % f, val, fmt=fmt, fill=VERDE, locked=False, wrap=(k == 'cong'))
        _nota(ws, f, nota, desde='C', hasta='F', chars=125.0)
    ws.merge_cells('B%d:E%d' % (M_PE['cong'], M_PE['cong']))
    ws.row_dimensions[M_PE['cong']].height = 32
    # C. coste por kg útil
    _f(ws, 'A%d' % PE_RC, 'C. COSTE POR KG ÚTIL Y POR PIEZA', b=True, fill=GRIS)
    for i in range(NPE):
        _f(ws, '%s%d' % (PC[i], PE_RC), '=%s5' % PC[i], b=True, fill=GRIS, wrap=True)
    _f(ws, 'F%d' % PE_RC, 'Notas', b=True, fill=GRIS)
    ws.row_dimensions[PE_RC].height = 62
    for k, rot, fmt, expr, nota in SAL_C:
        f = S_C[k]
        _f(ws, 'A%d' % f, rot, b=k == 'cku')
        for i in range(NPE):
            c = PC[i]
            ref = dict((kk, '$B$%d' % ff) for kk, ff in M_PE.items())
            ref.update((kk, '%s$%d' % (c, ff)) for kk, ff in P_ROW.items())
            ref.update((kk, '%s$%d' % (c, ff)) for kk, ff in S_C.items())
            _f(ws, '%s%d' % (c, f), '=' + a_excel(expr, ref), fmt=fmt, b=k == 'cku')
        _nota(ws, f, nota, desde='F', hasta='F', chars=75.0)
    # D. entero o lomo limpio
    _f(ws, 'A%d' % PE_RD, 'D. ¿ENTERO O LOMO LIMPIO?', b=True, fill=GRIS)
    _f(ws, 'B%d' % PE_RD, 'Valor', b=True, fill=GRIS)
    _f(ws, 'C%d' % PE_RD, 'Notas', b=True, fill=GRIS)
    ref = dict((k, '$B$%d' % f) for k, f in M_PE.items())
    ref.update((k, '$B$%d' % f) for k, f in S_D.items())
    ref.update({'rd0': '$B$%d' % P_ROW['rd'], 'cku0': '$B$%d' % S_C['cku'], 'cku1': '$C$%d' % S_C['cku'],
                'rem_r': '$B$%d:$B$%d' % (PE_O0, PE_O1), 'opc_r': '$A$%d:$A$%d' % (PE_O0, PE_O1)})
    for k, rot, fmt, expr, nota in SAL_D:
        f = S_D[k]
        neg = k in ('lab_kg', 'X', 'sem')
        _f(ws, 'A%d' % f, rot, b=neg, wrap=True)
        _f(ws, 'B%d' % f, '=' + a_excel(expr, ref), fmt=fmt, b=neg, wrap=(fmt == 'General'))
        _nota(ws, f, nota, desde='C', hasta='F', chars=125.0)
    ws.merge_cells('B%d:E%d' % (S_D['rem'], S_D['rem']))
    ws.row_dimensions[S_D['rem']].height = 48
    ws.row_dimensions[S_D['X']].height = 32
    ws.row_dimensions[S_D['sem']].height = 32
    _semaforo(ws, 'B%d' % S_D['sem'], SEM_ENT, SEM_LOMO)
    # lista del desplegable
    _f(ws, 'A%d' % PE_RO, 'LISTA DEL DESPLEGABLE DE CONGELACIÓN', b=True, fill=GRIS)
    _f(ws, 'B%d' % PE_RO, 'Recordatorio', b=True, fill=GRIS)
    for i, (op, rem) in enumerate(OPC):
        f = PE_O0 + i
        _f(ws, 'A%d' % f, op, wrap=True)
        _f(ws, 'B%d' % f, rem, wrap=True)
        ws.merge_cells('B%d:F%d' % (f, f))
        ws.row_dimensions[f].height = 32
    # validaciones
    _dv(ws, 'decimal', ['%s%d' % (c, P_ROW['pr']) for c in PC], f1='0', op='greaterThan', msg='Escribe un número mayor que 0.')
    _dv(ws, 'decimal', ['%s%d' % (c, P_ROW['rd']) for c in PC], f1='0.3', f2='1', op='between', titulo='Rendimiento no válido',
        msg='Escribe entre 0,3 y 1 (0,55 = 55 %).')
    _dv(ws, 'decimal', ['%s%d' % (c, P_ROW['rc']) for c in PC] + ['B%d' % M_PE['m']], f1='0', f2='0.5', op='between',
        titulo='Porcentaje no válido', msg='Escribe un porcentaje entre 0 y 0,5 (0,05 = 5 %).')
    _dv(ws, 'decimal', ['B%d' % M_PE[k] for k in ('gn', 'gs', 'gr', 'pe')], f1='0', op='greaterThan', msg='Escribe un número mayor que 0.')
    _dv(ws, 'decimal', ['B%d' % M_PE[k] for k in ('tl', 'hh')], f1='0', op='greaterThanOrEqual', msg='Escribe un número mayor o igual que 0.')
    _dv(ws, 'list', ['B%d' % M_PE['cong']], f1='$A$%d:$A$%d' % (PE_O0, PE_O1), titulo='Opción no válida', msg='Elige una opción de la lista.')
    _imprimir(ws, horizontal=True)


def hoja_arroz(wb):
    ws = wb.create_sheet(HA)
    for col, w in (('A', 60), ('B', 15), ('C', 15), ('D', 16), ('E', 14), ('F', 14), ('G', 14), ('H', 12), ('I', 15)):
        ws.column_dimensions[col].width = w
    _f(ws, 'A1', 'CALCULADORA DEL ARROZ — CUÁNTO SHARI SALE DE UN KILO, CUÁNTO CUESTA Y CUÁNTO COMPRAR AL DÍA', b=True, sz=14)
    _f(ws, 'A2', 'Escribe solo en las celdas VERDES: todo lo demás se calcula. Importes sin IVA. Todos los valores por defecto son una '
                 'estimación del kit, editable: pesa tus bolas de shari y cambia los gramos por los tuyos.', wrap=True)
    ws.merge_cells('A2:I2')
    ws.row_dimensions[2].height = 32
    _f(ws, 'A3', 'AVISO FIJO. El pH de cada tanda se registra en el Kit de Tareas Sushi Bar (plantilla 02). El shari a temperatura ambiente, sólo '
                 'si tu APPCC lo justifica con el pH de cada tanda y un tiempo máximo y lo demuestras a sanidad (RD 1086/2020, art. 30.5); si no, refrigerado.', wrap=True, b=True, fill=GRIS)
    ws.merge_cells('A3:I3')
    ws.row_dimensions[3].height = 36
    _bloque(ws, 'A. EL ARROZ Y EL SHARI', ENT_A, M_A, M_RA, SAL_A, S_A, ('kgs', 'cs', 'nigkg', 'arroz_dia'))
    cel = lambda k: 'B%d' % M_A[k]
    _dv(ws, 'decimal', [cel(k) for k in ('ak', 'zp', 'f', 'g_nig', 'g_hos', 'g_ura', 'g_tem', 'g_pok', 'gc')], f1='0', op='greaterThan',
        msg='Escribe un número mayor que 0.')
    _dv(ws, 'decimal', [cel(k) for k in ('z', 'nori', 'n_nig', 'n_hos', 'n_ura', 'n_tem', 'n_pok')], f1='0', op='greaterThanOrEqual',
        msg='Escribe un número mayor o igual que 0.')
    _dv(ws, 'decimal', [cel('mt')], f1='0', f2='0.5', op='between', titulo='Porcentaje no válido', msg='Escribe un porcentaje entre 0 y 0,5 (0,05 = 5 %).')
    _dv(ws, 'whole', [cel('cd')], f1='1', f2='1000000', op='between', titulo='Número no válido', msg='Escribe un número entero de 1 o más.')
    _imprimir(ws, horizontal=True)


def hoja_escandallo(wb):
    ws = wb.create_sheet(HE)
    for col, w in ((CN, 44), (CPF, 12), (CNE, 11), (CSH, 11), (CNO, 10), (COT, 13), (CPVP, 13), (CUD, 12), (CCO, 12), (CPS, 12), (CFC, 11),
                   (CSE, 17), (CSG, 14), (CMS, 12), (CML, 12), (CMP, 13), (CRK, 10)):
        ws.column_dimensions[col].width = w
    _f(ws, 'A1', 'ESCANDALLO DEL SUSHI BAR — 14 PRODUCTOS, POR PIEZA Y POR CANAL', b=True, sz=14)
    _f(ws, 'A2', 'Escribe solo en las celdas VERDES (nombres, pescado, gramos, nori, precios de carta y unidades). Importes sin IVA salvo el precio '
                 'de carta. Todos los valores por defecto son una estimación del kit, editable: pon los tuyos. El coste del pescado viene de la '
                 'hoja «%s», el del shari y el nori de la hoja «%s».' % (HP, HA), wrap=True)
    ws.merge_cells('A2:%s2' % CRK)
    ws.row_dimensions[2].height = 32
    for col, t in zip('ABD', ('DATOS GENERALES', 'Valor', 'Notas')):
        _f(ws, '%s4' % col, t, b=True, fill=GRIS)
    for k, rot, fmt, val, nota in ESC_IN:
        f = E[k]
        _f(ws, 'A%d' % f, rot)
        _f(ws, 'B%d' % f, val, fmt=fmt, fill=VERDE, locked=False)
        _nota(ws, f, nota, desde='D', hasta=CRK, chars=150.0)
    _f(ws, 'A%d' % E_SH, 'Coste por kg de shari con merma (€, de la hoja Arroz)')
    _f(ws, 'B%d' % E_SH, "='%s'!$B$%d" % (HA, S_A['csm']), fmt=EUR)
    _f(ws, 'A%d' % E_NO, 'Nori (€ por hoja, de la hoja Arroz)')
    _f(ws, 'B%d' % E_NO, "='%s'!$B$%d" % (HA, M_A['nori']), fmt=EUR3)
    cab = [(CN, 'PRODUCTO'), (CPF, 'Pescado (0 a 4)'), (CNE, 'Gramos de neta'), (CSH, 'Gramos de shari'), (CNO, 'Hojas de nori'),
           (COT, 'Otros ingredientes (€)'), (CPVP, 'PVP con IVA de sala (€)'), (CUD, 'Unidades a la semana'), (CCO, 'COSTE (€)'),
           (CPS, 'PVP sin IVA (€)'), (CFC, 'FOOD COST (%)'), (CSE, 'Semáforo'), (CSG, 'PVP con IVA sugerido (€)'),
           (CMS, 'Margen en sala (€)'), (CML, 'Margen para llevar (€)'), (CMP, 'Margen en plataforma (€)'), (CRK, 'Puesto por margen')]
    for col, t in cab:
        _f(ws, '%s%d' % (col, E_CAB), t, b=True, fill=GRIS, wrap=True)
    ws.row_dimensions[E_CAB].height = 48
    kg = "'%s'!$B$%d" % (HP, M_PE['KG'])
    fila = "'%s'!$%%s$%d" % (HP, S_C['cku'])
    elige = 'CHOOSE(%s%%d+1,0,%s)' % (CPF, ','.join(fila % c for c in PC))
    for j, (nom, pf, ne, sh, no, ot, pvp, ud) in enumerate(PRODS):
        f = E_P0 + j
        _f(ws, 'A%d' % f, nom, fill=VERDE, locked=False, b=True)
        _f(ws, '%s%d' % (CPF, f), pf, fmt='0', fill=VERDE, locked=False)
        _f(ws, '%s%d' % (CNE, f), float(ne), fmt='#,##0.0;-#,##0.0;"-"', fill=VERDE, locked=False)
        _f(ws, '%s%d' % (CSH, f), float(sh), fmt='#,##0.0;-#,##0.0;"-"', fill=VERDE, locked=False)
        _f(ws, '%s%d' % (CNO, f), float(no), fmt='#,##0.00;-#,##0.00;"-"', fill=VERDE, locked=False)
        _f(ws, '%s%d' % (COT, f), float(ot), fmt=EUR, fill=VERDE, locked=False)
        _f(ws, '%s%d' % (CPVP, f), float(pvp), fmt=EUR, fill=VERDE, locked=False, b=True)
        _f(ws, '%s%d' % (CUD, f), float(ud), fmt=ENT, fill=VERDE, locked=False)
        ex = {
            CCO: '%s%d/%s*%s+%s%d/%s*$B$%d+%s%d*$B$%d+%s%d' % (CNE, f, kg, elige % f, CSH, f, kg, E_SH, CNO, f, E_NO, COT, f),
            CPS: '%s%d/(1+$B$%d)' % (CPVP, f, E['iva']),
            CFC: '%s%d/%s%d' % (CCO, f, CPS, f),
            CSE: 'IF(%s%d<=$B$%d,"%s",IF(%s%d<=$B$%d+$B$%d,"%s","%s"))' % (
                CFC, f, E['obj'], FC_OK, CFC, f, E['obj'], E['tol'], FC_VIG, FC_ALTO),
            CSG: '%s%d/$B$%d*(1+$B$%d)' % (CCO, f, E['obj'], E['iva']),
            CMS: '%s%d-%s%d' % (CPS, f, CCO, f),
            CML: '%s%d-$B$%d' % (CMS, f, E['cl']),
            CMP: '%s%d*(1+$B$%d)*(1-$B$%d)-%s%d-$B$%d' % (CPS, f, E['rec'], E['com'], CCO, f, E['cp']),
            CRK: '1+COUNTIF($%s$%d:$%s$%d,">"&%s%d)' % (CMS, E_P0, CMS, E_P1, CMS, f),
        }
        for _k, col, fmt in SALP:
            neg = col in (CCO, CFC)
            _f(ws, '%s%d' % (col, f), '=' + ex[col], fmt=fmt, b=neg or col == CRK, wrap=(col == CSE))
        ws.row_dimensions[f].height = 32
    rng = '%s%d:%s%d' % (CSE, E_P0, CSE, E_P1)
    top = '%s%d' % (CSE, E_P0)
    for txt, fill in ((FC_OK, OK), (FC_VIG, AMBAR), (FC_ALTO, ROJO)):
        ws.conditional_formatting.add(rng, FormulaRule(formula=['%s="%s"' % (top, txt)], fill=fill))
    # ponderado
    _f(ws, 'A%d' % E_T0, 'FOOD COST PONDERADO DE LA SEMANA (suma de costes × unidades ÷ suma de PVP sin IVA × unidades)', b=True, fill=GRIS)
    _f(ws, 'B%d' % E_T0, 'Valor', b=True, fill=GRIS)
    rc = '$%s$%d:$%s$%d' % (CCO, E_P0, CCO, E_P1)
    rp = '$%s$%d:$%s$%d' % (CPS, E_P0, CPS, E_P1)
    ru = '$%s$%d:$%s$%d' % (CUD, E_P0, CUD, E_P1)
    exs = {'ing': 'SUMPRODUCT(%s,%s)' % (rp, ru), 'cos': 'SUMPRODUCT(%s,%s)' % (rc, ru),
           'fcp': 'IF($B$%d=0,0,$B$%d/$B$%d)' % (T['ing'], T['cos'], T['ing']),
           'semp': 'IF($B$%d<=$B$%d,"%s",IF($B$%d<=$B$%d+$B$%d,"%s","%s"))' % (
               T['fcp'], E['obj'], FC_OK, T['fcp'], E['obj'], E['tol'], FC_VIG, FC_ALTO)}
    for k, rot, fmt, _x in TOT:
        f = T[k]
        _f(ws, 'A%d' % f, rot, b=k == 'fcp')
        _f(ws, 'B%d' % f, '=' + exs[k], fmt=fmt, b=k == 'fcp', wrap=(k == 'semp'))
    ws.row_dimensions[T['semp']].height = 32
    for txt, fill in ((FC_OK, OK), (FC_VIG, AMBAR), (FC_ALTO, ROJO)):
        ws.conditional_formatting.add('B%d' % T['semp'], FormulaRule(formula=['$B$%d="%s"' % (T['semp'], txt)], fill=fill))
    _f(ws, 'A%d' % E_FIN, 'Pescado: 0 = ninguno · 1 = %s · 2 = %s · 3 = %s · 4 = %s. El food cost de la tabla es el del plato servido en sala, sin '
                          'contar las bebidas. En el pedido para llevar se descuenta la caja; en la plataforma, además, se sube el precio con el '
                          'recargo y se descuenta la comisión sobre ese precio. El margen se ordena por el de sala. '
                          'AVISO FIJO: la tempura y las gyozas viajan mal; en la plataforma, ofrécelas solo con envase ventilado o déjalas fuera '
                          'de la carta de reparto. Este coste es TEÓRICO: el plan de negocio trabaja con un 29 %%, el lado prudente; la diferencia '
                          'entre la receta y la caja está explicada en la hoja «Instrucciones».' % tuple(p[0] for p in PESCADOS),
       i=True, wrap=True)
    ws.merge_cells('A%d:%s%d' % (E_FIN, CRK, E_FIN))
    ws.row_dimensions[E_FIN].height = 82
    ws.freeze_panes = 'B%d' % (E_CAB + 1)
    # validaciones
    _dv(ws, 'decimal', ['B%d' % E[k] for k in ('iva', 'tol', 'rec', 'com')], f1='0', f2='0.9', op='between',
        titulo='Porcentaje no válido', msg='Escribe un porcentaje entre 0 y 0,9 (0,1 = 10 %).')
    _dv(ws, 'decimal', ['B%d' % E['obj']], f1='0.01', f2='0.9', op='between', titulo='Objetivo no válido',
        msg='El objetivo no puede ser 0 (se divide entre él): escribe entre 0,01 y 0,9.')
    _dv(ws, 'decimal', ['B%d' % E['cl'], 'B%d' % E['cp']]
        + ['%s%d' % (c, f) for c in (CNE, CSH, CNO, COT, CPVP, CUD) for f in range(E_P0, E_P1 + 1)],
        f1='0', op='greaterThanOrEqual', titulo='Importe no válido', msg='Escribe un número mayor o igual que 0.')
    _dv(ws, 'whole', ['%s%d' % (CPF, f) for f in range(E_P0, E_P1 + 1)], f1='0', f2=str(NPE), op='between', titulo='Pescado no válido',
        msg='Escribe 0 (ninguno), 1, 2, 3 o 4, el número de la columna de la hoja Pescado.')
    _imprimir(ws, horizontal=True)


INSTRUCCIONES = [
    ('INSTRUCCIONES DE USO — Calculadora del pescado, el arroz y el escandallo', True),
    ('Tres herramientas en un libro: la hoja «Pescado» te dice cuánto cuesta el kilo útil de salmón entero, de lomo comprado limpio y de atún, '
     'cuánto vale cada loncha, corte y roll y si te compensa comprar el pescado entero; la hoja «Arroz» calcula el shari que sale de un kilo de '
     'arroz crudo, su coste, cuántos nigiris das por kilo y cuánto arroz compras al día; la hoja «Escandallo» calcula el coste, el food cost, '
     'el margen y el precio de carta de 14 productos, con el margen en sala, para llevar y en plataforma, y el food cost ponderado de tu '
     'semana. Solo se escribe en las celdas VERDES; lo demás se calcula.', False),
    ('Todos los valores por defecto son estimación del kit, editable: son ejemplos coherentes con el caso del plan de negocio, no precios '
     'de nadie. Cámbialos por los de tu proveedor y por lo que midas en tu barra antes de decidir precios.', False),
    ('HOJA «PESCADO»', True),
    ('1. Mide tu rendimiento. Pesa 3 salmones enteros, límpialos y pesa el lomo que sale de cada uno; divide entre el peso del pescado entero '
     'y escribe la media en «Rendimiento». Pesa también el recorte que reaprovechas (tartar, maki, poke) y escríbelo en «Recorte que '
     'reaprovechas». Pesa 10 lonchas de nigiri y 10 cortes de sashimi y usa la media en los gramos.', False),
    ('2. El coste por kilo útil es el precio ÷ (rendimiento + recorte) ÷ (1 − merma de vitrina). Es el número que alimenta el escandallo: '
     'con los valores del kit sale a 14,29 €/kg en el salmón entero, 24,16 en el lomo comprado y 28,81 en el atún en saku.', False),
    ('3. La limpieza también cuesta. El kit supone 25 minutos por pieza entera, un dato sin fuente: cronométralo. El coste por hora del '
     'itamae (el jefe de barra: compra, limpia y corta el pescado) se estima en 22 €/h: 2.200 € brutos × 14 pagas + 33 % de cargas ÷ 1.840 h '
     '≈ 22 €/h. La hoja te dice hasta qué coste de limpieza por kilo útil te compensa comprar el pescado entero.', False),
    ('4. Anisakis (RD 1021/2022, art. 8). El pescado que se sirve crudo o casi crudo se congela a −20 °C o menos durante al menos 24 horas, o a '
     '−35 °C o menos durante al menos 15 horas, en la totalidad del producto. La acuicultura queda exenta solo con documento escrito del '
     'proveedor. Hay que informar al cliente en la carta o en un cartel (art. 8.2). El desplegable de la hoja es solo un recordatorio: no '
     'cambia ningún cálculo.', False),
    ('HOJA «ARROZ»', True),
    ('1. Pesa 10 bolas de shari (el arroz ya aliñado con sushi-zu) y usa la media en los gramos por nigiri; haz lo mismo con un hosomaki, un '
     'uramaki, un temaki y un poke.', False),
    ('2. Los resultados clave son los kilos de shari por kilo de arroz crudo (2,42 con los valores del kit), el coste por kilo de shari y los '
     'nigiris que das por cada kilo de arroz crudo. Con eso sabes cuánto arroz tienes que comprar al día para tus comensales.', False),
    ('3. Seguridad del shari. A temperatura ambiente solo puede mantenerse si tu APPCC lo justifica y lo demuestras a sanidad (RD 1086/2020, art. 30.5); '
     'si no, se mantiene refrigerado. El pH de 4,6 o menos es el límite crítico del APPCC del Kit de Tareas Sushi Bar, no una cifra legal. '
     'El pH de cada tanda se registra en el Kit de Tareas Sushi Bar (plantilla 02).', False),
    ('HOJA «ESCANDALLO»', True),
    ('1. Escribe el IVA, el food cost objetivo, el margen de vigilancia, las cajas de para llevar y de plataforma, el recargo y la comisión de la '
     'plataforma.', False),
    ('2. En la tabla, cada fila es un producto. «Pescado» es el número de la columna de la hoja Pescado que usa (1 salmón entero, 2 lomo '
     'comprado, 3 atún, 4 libre; 0 si no lleva). Escribe los gramos de neta, los gramos de shari, las hojas de nori y lo que cuestan el resto de '
     'ingredientes, el precio de carta con IVA y las unidades que vendes a la semana.', False),
    ('3. Para cada producto ves el coste, el precio sin IVA, el food cost con semáforo, el precio de carta con IVA que necesitarías para llegar '
     'a tu objetivo y el margen en sala, para llevar y en plataforma, con el puesto de cada producto por margen en sala. Abajo, el food cost '
     'ponderado de tu semana. La tempura y las gyozas viajan mal: vigílalas en la plataforma.', False),
    ('DEL FOOD COST TEÓRICO AL 29 % DEL PLAN', True),
    ('El food cost de esta hoja es teórico: el que saldría si todo lo que compras se vendiera tal como está en el escandallo. Con los '
     'valores del kit, la semana ponderada sale en torno a un 16 %. El plan de negocio trabaja con un 29 %, el lado prudente: entre la receta y '
     'la caja se pierden la merma del pescado que no se vende en la vitrina y del recorte que no se reaprovecha, la neta y el shari que pesan '
     'más que en el escandallo, el shari, el nori y las salsas que se tiran al cierre, las comidas del personal, las invitaciones y los errores '
     'de pedido, y el menú del día, que deja poco margen; además, los precios del pescado son estimaciones del kit. Un restaurante japonés con '
     'pescado de calidad suele trabajar con un 30-35 %.', False),
    ('Tu comparación real es esta: las compras de comida del mes ÷ las ventas de comida sin IVA del mes. Si después del primer trimestre sale '
     'por encima o por debajo del 29 %, cambia la celda «Coste de mercancía sobre las ventas de COMIDA» del plan: cada punto vale unos 4.000 € al año.',
     False),
    ('PRECIOS, IVA Y FUENTES', True),
    ('El IVA de la comida preparada para consumir en el local, para llevar o a domicilio es el 10 % (Ley 37/1992, art. 91). La compra de arroz '
     'lleva el 4 %. Los importes de la calculadora van sin IVA, salvo el precio de carta.', False),
    ('Los precios y rendimientos salen del estudio de mercado del kit (octubre de 2026): salmón y arroz de mayorista y minorista, nori de '
     'tienda especializada y una carta de cadena de sushi como referencia de precios. Ninguna cifra de esta calculadora es un dato publicado '
     'de tu negocio: sustitúyelas por las tuyas.', False),
]


def hoja_instrucciones(wb):
    ws = wb.create_sheet(HI)
    ws.column_dimensions['A'].width = 110
    f = 1
    for t, b in INSTRUCCIONES:
        _f(ws, 'A%d' % f, t, b=b, sz=13 if f == 1 else 11, wrap=True)
        if not b:
            ws.row_dimensions[f].height = max(16, 15 * math.ceil(len(t) / 120.0))
        f += 2 if b and f > 1 else 1
    f += 1
    for t in ('Para editar una celda que no esté en verde: Revisar → Desproteger hoja (no tiene contraseña).',
              'Más plantillas y kits del catálogo en aichef.pro/productos-digitales', motor.BIO_LINE,
              motor.VERSION_FMT.format(pid=PID)):
        _f(ws, 'A%d' % f, t, wrap=True)
        f += 1
    _imprimir(ws)


def generar(destino):
    wb = openpyxl.Workbook()
    hoja_pescado(wb)
    hoja_arroz(wb)
    hoja_escandallo(wb)
    hoja_instrucciones(wb)
    pr = wb.properties
    pr.title = '%s · v%s' % (TITULO, motor.VERSION)
    pr.subject = 'Plan de Negocio: Sushi Bar'
    pr.description = 'aichef.pro/%s' % PID
    pr.keywords = 'plan de negocio sushi bar, calculadora de pescado y arroz para sushi, escandallo de sushi, AI Chef Pro'
    pr.creator = pr.lastModifiedBy = motor.CREATOR
    os.makedirs(os.path.dirname(destino), exist_ok=True)
    wb.save(destino)
    sys.path.insert(0, PD)
    import inject_cache as IC                                                      # noqa: E402
    nf, inj, fail = IC.inject(destino)
    print('%s: formulas=%d cache_inyectado=%d fallos_pycel=%d' % (os.path.basename(destino), nf, inj, fail))
    return fail


def valores_expuestos(ruta):
    """Lee de la CACHÉ del xlsx los valores expuestos: {nombre: valor}."""
    wv = openpyxl.load_workbook(ruta, data_only=True)
    return OrderedDict((n, wv[h][c].value) for n, (h, c) in EXPUESTOS.items())


# ---------------------------------------------------------------------------------------------------------------------
# Autotest
# ---------------------------------------------------------------------------------------------------------------------
PROHIBIDOS = [r'pizz', r'napol[ei]t', r'taquer', r'trompo', r'\btacos?\b', r'tortilla', r'\bcomal\b', r'al pastor', r'gringa', r'nixtamal',
              r'mezcal', r'tequila', r'aguas? frescas?', r'carnitas', r'suadero', r'cochinita', r'guacamole', r'salsa macha',
              r'plan-negocio-taqueria', r'plan-negocio-polleria', r'mozzarella', r'poller', r'asador', r'\bpollo', r'broaster',
              r'sin gluten', r'auténtic', r'mexican', r'digital obligatorio', r'sales tax|service charge',
              r'1420/2006', r'7 días', r'sashimi grade|sushi grade',
              r'anisakis[^.]{0,80}[−-]\s?18|[−-]\s?18\s?°[^.]{0,80}anisakis',
              r'obligatori[^.]{0,60}\bpH\b|\bpH\b[^.]{0,60}obligatori|obligatori[^.]{0,40}4 horas|4 horas[^.]{0,40}obligatori']
CELDA = {}                                                                         # clave de entrada → (hoja, celda)
for _i in range(NPE):
    for _k in ('pr', 'rd', 'rc'):
        CELDA['%s%d' % (_k, _i)] = (HP, '%s%d' % (PC[_i], P_ROW[_k]))
for _k, _f2 in M_PE.items():
    CELDA[_k] = (HP, 'B%d' % _f2)
for _k, _f2 in M_A.items():
    CELDA[_k] = (HA, 'B%d' % _f2)
for _k, _f2 in E.items():
    CELDA[_k] = (HE, 'B%d' % _f2)
for _j in range(NPR):
    for _k, _c in (('pf', CPF), ('ne', CNE), ('sh', CSH), ('no', CNO), ('ot', COT), ('pvp', CPVP), ('ud', CUD)):
        CELDA['%s%d' % (_k, _j)] = (HE, '%s%d' % (_c, E_P0 + _j))


def celdas_salida(r):
    """(hoja, celda, valor esperado, etiqueta) de todas las salidas."""
    out = []
    for i, c in enumerate(r['Pe']):
        for k in S_C:
            out.append((HP, '%s%d' % (PC[i], S_C[k]), c[k], 'Pe%d.%s' % (i, k)))
    for k in S_D:
        out.append((HP, 'B%d' % S_D[k], r['D'][k], 'D.' + k))
    for k in S_A:
        out.append((HA, 'B%d' % S_A[k], r['A'][k], 'A.' + k))
    out.append((HE, 'B%d' % E_SH, r['shari'], 'shari'))
    out.append((HE, 'B%d' % E_NO, r['nori'], 'nori'))
    for j, q in enumerate(r['Pr']):
        for k, col, _fm in SALP:
            out.append((HE, '%s%d' % (col, E_P0 + j), q[k], '%s[%d]' % (k, j)))
    for k in T:
        out.append((HE, 'B%d' % T[k], r['T'][k], 'T.' + k))
    return out


def _distinto(x, v):
    if isinstance(v, str):
        return x != v
    return not isinstance(x, (int, float)) or abs(float(x) - v) > 0.01


def _checks_logica(nombre, jj, err):
    r = calcular(jj)
    for i, c in enumerate(r['Pe']):
        if abs(c['cku'] * c['tot'] * (1 - jj['m']) - jj['pr%d' % i]) > 1e-6:
            err.append('%s: coste útil × (r + rec) × (1 − merma) ≠ precio en la columna %d' % (nombre, i))
        if abs(c['c_nig'] * jj['KG'] / jj['gn'] - c['cku']) > 1e-6:
            err.append('%s: coste de loncha ≠ coste útil × gramos ÷ 1.000 en la columna %d' % (nombre, i))
    a = r['A']
    if abs(a['nigkg'] - 1000 * jj['f'] * (1 + jj['z'] / 1000) / jj['g_nig']) > 1e-6:
        err.append('%s: nigiris por kg ≠ 1.000 × f × (1 + z ÷ 1.000) ÷ g' % nombre)
    if abs(a['cs'] * a['kgs'] - (jj['ak'] + jj['f'] * jj['z'] / 1000 * jj['zp'])) > 1e-6:
        err.append('%s: coste por kg de shari × kg de shari ≠ coste del arroz y el sushi-zu' % nombre)
    if abs(a['arroz_dia'] * a['kgs'] * 1000 - jj['cd'] * jj['gc']) > 1e-6:
        err.append('%s: arroz al día × shari por kg ≠ comensales × gramos' % nombre)
    for j in (J_C24, 9):
        q = r['Pr'][j]
        if not q['plat'] < q['llevar'] < q['sala']:
            err.append('%s: margen plataforma < para llevar < sala no se cumple en el combinado %d' % (nombre, j))
    for j, q in enumerate(r['Pr']):
        if not q['plat'] < q['llevar'] < q['sala']:
            err.append('%s: margen plataforma < para llevar < sala no se cumple en el producto %d' % (nombre, j))
    if abs(r['D']['X'] - (r['Pe'][1]['cku'] - r['Pe'][0]['cku'])) > 1e-9:
        err.append('%s: X ≠ coste útil del lomo − coste útil del entero' % nombre)


def autotest(destino):
    from pycel import ExcelCompiler
    err, n = [], 0
    wv = openpyxl.load_workbook(destino, data_only=True)
    wf = openpyxl.load_workbook(destino)
    base = valores_defecto()
    j2 = dict(base, pr0=11.95, rd0=0.5, rc0=0.1, pr1=20.7, rc1=0.0, pr2=30.0, rd2=0.85, pr3=12.0, rd3=0.7, m=0.08, gn=18.0, gs=20.0, gr=50.0,
              tl=35.0, hh=18.0, pe=4.0, cong=OPC[1][0], ak=2.5, f=2.0, z=80, zp=6.0, g_nig=22.0, g_ura=100.0, nori=0.3, n_ura=0.6, mt=0.1,
              cd=100, gc=170.0, iva=0.10, obj=0.28, tol=0.04, cl=0.5, cp=1.0, rec=0.1, com=0.25, pf3=2, ne3=50, sh6=100.0, ot9=0.5,
              pvp10=31.0, ud0=100, ud10=40)
    j3 = dict(base, pr0=8.0, rd0=0.6, pr1=23.0, rd1=0.9, pr2=25.0, pr3=40.0, rd3=0.6, rc3=0.1, m=0.02, gn=12.0, tl=15.0, hh=16.0, pe=6.0,
              cong=OPC[2][0], ak=1.7, f=2.5, z=120, zp=3.5, g_nig=18.0, nori=0.2, mt=0.0, cd=60, gc=130.0, obj=0.35, tol=0.1, cl=0.3, cp=0.8,
              rec=0.2, com=0.35, pf0=4, ne0=40, pf10=2, ne10=200, sh10=300.0, pvp9=16.0, pvp10=26.0, ud5=0, ud12=200)
    # 1. caché y pycel del juego por defecto; pycel de los otros dos con set_value
    xl = ExcelCompiler(filename=destino)
    r0 = calcular(base)
    for hoja, c, v, lab in celdas_salida(r0):
        n += 1
        for et, x in (('pycel', xl.evaluate("'%s'!%s" % (hoja, c))), ('caché', wv[hoja][c].value)):
            if _distinto(x, v):
                err.append('defecto · %s!%s (%s): %s %r ≠ esperado %r' % (hoja, c, lab, et, x, v))
    for nombre, jj in (('juego 2', j2), ('juego 3', j3)):
        r = calcular(jj)
        xl2 = ExcelCompiler(filename=destino)
        for hoja, c, v, lab in celdas_salida(r0):
            xl2.evaluate("'%s'!%s" % (hoja, c))
        for k, v in jj.items():
            if v != base[k]:
                h, c = CELDA[k]
                xl2.set_value("'%s'!%s" % (h, c), v)
        for hoja, c, v, lab in celdas_salida(r):
            n += 1
            x = xl2.evaluate("'%s'!%s" % (hoja, c))
            if _distinto(x, v):
                err.append('%s · %s!%s (%s): pycel %r ≠ esperado %r' % (nombre, hoja, c, lab, x, v))
    # 2. identidades de la lógica en los tres juegos
    for nombre, jj in (('defecto', base), ('juego 2', j2), ('juego 3', j3)):
        n += 12
        _checks_logica(nombre, jj, err)
    # identidades sobre la hoja viva (pycel, valores por defecto)
    ev = lambda h, c: xl.evaluate("'%s'!%s" % (h, c))
    for i in range(NPE):
        n += 1
        pr_, rd_, rc_ = (ev(HP, '%s%d' % (PC[i], P_ROW[k])) for k in ('pr', 'rd', 'rc'))
        m_ = ev(HP, 'B%d' % M_PE['m'])
        if abs(ev(HP, '%s%d' % (PC[i], S_C['cku'])) - pr_ / (rd_ + rc_) / (1 - m_)) > 1e-6:
            err.append('hoja Pescado, columna %s: coste útil ≠ precio ÷ (r + rec) ÷ (1 − m)' % PC[i])
    kgs, ng, gn = ev(HA, 'B%d' % S_A['kgs']), ev(HA, 'B%d' % S_A['nigkg']), ev(HA, 'B%d' % M_A['g_nig'])
    fz, zz = ev(HA, 'B%d' % M_A['f']), ev(HA, 'B%d' % M_A['z'])
    n += 1
    if abs(ng - 1000 * fz * (1 + zz / 1000) / gn) > 1e-6 or abs(kgs * 1000 / gn - ng) > 1e-6:
        err.append('hoja Arroz: nigiris por kg ≠ 1.000 × f × (1 + z ÷ 1.000) ÷ g')
    # 3. valores con sentido
    nig, c24 = r0['Pr'][J_NIG], r0['Pr'][J_C24]
    if not 0.08 <= nig['fc'] <= 0.15:
        err.append('food cost del nigiri de salmón %.1f %% fuera de 8-15 %%' % (nig['fc'] * 100))
    if not 0.12 <= c24['fc'] <= 0.20:
        err.append('food cost del combinado de 24 %.1f %% fuera de 12-20 %%' % (c24['fc'] * 100))
    if not 0.14 <= r0['T']['fcp'] <= 0.23:
        err.append('food cost ponderado por defecto %.1f %% fuera de 14-23 %%' % (r0['T']['fcp'] * 100))
    ck = [c['cku'] for c in r0['Pe'][:3]]
    if any(abs(a_ - b_) > 0.01 for a_, b_ in zip(ck, (14.29, 24.15, 28.81))):
        err.append('coste por kg útil por defecto %r ≠ 14,29 · 24,15 · 28,81' % [round(x, 3) for x in ck])
    A_ = r0['A']
    if abs(A_['kgs'] - 2.42) > 0.005 or abs(A_['cs'] - 1.16) > 0.01 or abs(A_['nigkg'] - 121) > 0.5 or abs(A_['arroz_dia'] - 4.9) > 0.05:
        err.append('arroz por defecto: 2,42 · 1,16 · 121 · 4,9 (salió %r)' % [round(A_[k], 3) for k in ('kgs', 'cs', 'nigkg', 'arroz_dia')])
    if r0['D']['sem'] != SEM_ENT or abs(r0['D']['X'] - 9.87) > 0.02:
        err.append('entero o lomo por defecto: «%s», X = %.3f (esperado entero y 9,87)' % (r0['D']['sem'], r0['D']['X']))
    if calcular(dict(base, tl=240.0))['D']['sem'] != SEM_LOMO:
        err.append('el veredicto no pasa a «lomo limpio» con 240 minutos de limpieza')
    if sorted(q['rank'] for q in r0['Pr']) != list(range(1, NPR + 1)) and len(set(q['rank'] for q in r0['Pr'])) < NPR - 2:
        err.append('los puestos por margen no son coherentes')
    # 4. valores expuestos: caché frente a calcular()
    esperados = {'COSTE_KG_UTIL_SALMON': r0['Pe'][0]['cku'], 'COSTE_KG_UTIL_LOMO': r0['Pe'][1]['cku'], 'COSTE_KG_UTIL_ATUN': r0['Pe'][2]['cku'],
                 'KG_SHARI_KG_ARROZ': A_['kgs'], 'NIGIRIS_KG_ARROZ': A_['nigkg'], 'ARROZ_DIA_KG': A_['arroz_dia'],
                 'FC_NIGIRI_SALMON': nig['fc'], 'FC_COMBINADO_24': c24['fc'], 'MARGEN_SALA_COMB24': c24['sala'],
                 'MARGEN_PLATAFORMA_COMB24': c24['plat'], 'FC_PONDERADO': r0['T']['fcp']}
    expuestos = valores_expuestos(destino)
    for nombre, e_ in esperados.items():
        n += 1
        if _distinto(expuestos[nombre], e_):
            err.append('expuesto %s: caché %r ≠ %r' % (nombre, expuestos[nombre], e_))
    # 5. formato, protección, DV, fórmulas, textos
    if motor.gate_formatos(wf, FICHERO):
        err.append('gate_formatos: %r' % motor.gate_formatos(wf, FICHERO)[:3])
    dvh = dict((w.title, dict((str(rng), dv) for dv in w.data_validations.dataValidation for rng in dv.sqref.ranges))
               for w in (wf[HP], wf[HA], wf[HE]))
    for hoja, k, tipo in ((HP, 'pr0', 'decimal'), (HP, 'rd1', 'decimal'), (HP, 'm', 'decimal'), (HP, 'cong', 'list'), (HA, 'ak', 'decimal'),
                          (HA, 'cd', 'whole'), (HE, 'obj', 'decimal'), (HE, 'iva', 'decimal'), (HE, 'com', 'decimal'), (HE, 'pf0', 'whole'),
                          (HE, 'ud0', 'decimal')):
        c = CELDA[k][1]
        dv = dvh[hoja].get(c)
        if dv is None or dv.type != tipo:
            err.append('DV ausente o de otro tipo en %s!%s (esperada %s)' % (hoja, c, tipo))
    textos, todo_celdas = [], []
    for w in wf.worksheets:
        if w.page_setup.paperSize != 9 or not w.protection.sheet:
            err.append('%s: A4 y protección' % w.title)
        for row in w.iter_rows():
            for c in row:
                v = c.value
                if v is not None:
                    todo_celdas.append(str(v))
                if isinstance(v, str) and c.data_type != 'f':
                    textos.append(v)
                if c.data_type == 'f':
                    if wv[w.title][c.coordinate].value is None:
                        err.append('%s!%s: fórmula sin caché' % (w.title, c.coordinate))
                    if len(str(c.value)) >= 200:
                        err.append('%s!%s: fórmula de %d caracteres (≥ 200)' % (w.title, c.coordinate, len(str(c.value))))
                    sin_refs = re.sub(r'"[^"]*"', '', re.sub(r"'[^']+'!|\$?[A-Z]{1,2}\$?\d+", '', str(c.value)))
                    for num in re.findall(r'(?<![A-Za-z])\d+\.?\d*', sin_refs):
                        if num not in ('0', '1'):
                            err.append('%s!%s: número tecleado %s en %s' % (w.title, c.coordinate, num, c.value))
                if c.fill.fgColor.rgb in ('00E8F5E9', 'FFE8F5E9') and c.protection.locked:
                    err.append('%s!%s: verde bloqueada' % (w.title, c.coordinate))
                if c.fill.fgColor.rgb in ('00E8F5E9', 'FFE8F5E9') and c.data_type == 'f':
                    err.append('%s!%s: verde con fórmula' % (w.title, c.coordinate))
    for hoja, ks in ((HP, ('KG', 'MH')), (HA, ('KG', 'ML'))):
        for k in ks:
            c = wf[hoja]['B%d' % (M_PE if hoja == HP else M_A)[k]]
            if c.fill.fgColor.rgb not in ('00F2F2F2', 'FFF2F2F2') or c.protection.locked is not True:
                err.append('constante gris %s mal marcada o desbloqueada' % k)
    todo = '\n'.join(textos)
    barrido = '\n'.join(todo_celdas)
    for rx in PROHIBIDOS:
        if re.search(rx, barrido, re.I):
            err.append('prohibido /%s/' % rx)
    for m_ in re.finditer(r'art\. 30', barrido):
        if 'RD 1086/2020' not in barrido[max(0, m_.start() - 60):m_.end() + 60]:
            err.append('«art. 30» sin «RD 1086/2020» al lado: %r' % barrido[max(0, m_.start() - 40):m_.end() + 20])
    if re.search(r'[^\x00-ɏ‐-‧€−→·≥≈]', barrido):
        err.append('caracteres no latinos o raros: %r' % sorted(set(re.findall(r'[^\x00-ɏ‐-‧€−→·≥≈]', barrido))))
    if 'estimación del kit, editable' not in barrido.lower():
        err.append('falta el rótulo «estimación del kit, editable»')
    if 'sin fuente: mídelo' not in barrido.lower():
        err.append('falta «sin fuente: mídelo» en los minutos de limpieza')
    if motor.VERSION_FMT.format(pid=PID) not in todo or motor.BIO_LINE not in todo:
        err.append('falta la línea de versión o la bio')
    if '1.0' not in motor.VERSION_FMT or 'noviembre 2026' not in motor.VERSION_FMT:
        err.append('versión no es 1.0 · noviembre 2026: %s' % motor.VERSION_FMT)
    if wf.sheetnames != [HP, HA, HE, HI]:
        err.append('pestañas %r' % wf.sheetnames)
    if NPR != 14:
        err.append('no son 14 productos')
    p = wf.properties
    if (p.creator != 'AI Chef Pro' or p.description != 'aichef.pro/%s' % PID or 'AI Chef Pro' not in (p.keywords or '')
            or p.subject != 'Plan de Negocio: Sushi Bar'):
        err.append('docProps: creator=%r description=%r keywords=%r subject=%r' % (p.creator, p.description, p.keywords, p.subject))
    for x in err:
        print('ERROR', x)
    print('autotest calculadora sushi bar ES: %d comprobaciones (pycel + caché por defecto; pycel en 2 juegos más; identidades) · %s'
          % (n, 'VERDE' if not err else 'ROJO (%d errores)' % len(err)))
    print('  pescado (€/kg útil): ' + ' · '.join('%s %.2f' % (PESCADOS[i][0], c['cku']) for i, c in enumerate(r0['Pe'])))
    d = r0['D']
    print('  limpieza: %.3f €/kg útil · el entero compensa hasta %.3f €/kg útil · %s' % (d['lab_kg'], d['X'], d['sem']))
    print('  arroz: %.3f kg shari/kg crudo · %.3f €/kg shari (%.3f con merma) · %.1f nigiris/kg crudo · %.2f kg crudo/día'
          % (A_['kgs'], A_['cs'], A_['csm'], A_['nigkg'], A_['arroz_dia']))
    for nombre, (h, c) in EXPUESTOS.items():
        print('  %-26s %s!%-5s %s' % (nombre, h, c, expuestos[nombre]))
    for j, p_ in enumerate(PRODS):
        q = r0['Pr'][j]
        print('  %-46s coste %.3f € · food cost %5.1f %% · %s · margen sala %.2f / llevar %.2f / plataforma %.2f € · puesto %d · %d uds'
              % (p_[0], q['c_tot'], q['fc'] * 100, q['sem'], q['sala'], q['llevar'], q['plat'], q['rank'], p_[7]))
    print('  ponderado: ventas %.2f € · coste %.2f € · food cost %.1f %% · %s' % (r0['T']['ing'], r0['T']['cos'], r0['T']['fcp'] * 100, r0['T']['semp']))
    return 1 if err else 0


if __name__ == '__main__':
    ap = argparse.ArgumentParser()
    ap.add_argument('--salida', required=True, help='carpeta de salida')
    ap.add_argument('--autotest', action='store_true')
    a = ap.parse_args()
    destino = os.path.join(a.salida, FICHERO)
    if not a.autotest and generar(destino):
        sys.exit('ERROR: fallos de pycel en inject_cache')
    sys.exit(autotest(destino))
