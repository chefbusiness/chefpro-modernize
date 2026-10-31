#!/usr/bin/env python3
"""
semilla_sus.py — Plan de Negocio: Sushi Bar (ES 1.0; F1-camino §2-§3): prepara el ORIGEN que `planes-v2_0/main.py`
convierte en el plan financiero y el checklist del sushi bar. Copia adaptada de `plan-polleria/semilla_pol.py` SIN
`_sin_rotacion` (el sushi bar tiene sala que rota: la rotación de la Taquería se queda ENCENDIDA, F1-camino §0.1).

Copia los dos xlsx de la Taquería ES 1.0 publicados (`dl/plan-negocio-taqueria/`, SOLO plan y checklist, SOLO
LECTURA; molde A-β + checklist C2) con los nombres de fichero del sushi bar y escribe lo que el módulo de contenido no
alcanza:
  · plan: valores de `0. Supuestos` (grupo_a conserva lo verde que trae el libro; `B60`/`B61` ya son 7/12 en la
    Taquería: assert, sin escribir), títulos de hoja (Inversión A1/A2/A3, PyG A2, Punto Equilibrio A2), las 18
    partidas de inversión (filas 7-24: rótulo, importe, nota, ¿IVA?), los 9 costes fijos propios del P&L (filas
    26-34: rótulo e importe; la nota la pone el contenido) y los docProps description/keywords. `PyG!G13` ya es VALOR
    en la Taquería y lo rehace el motor con `IVA_COMPRAS` (0,087): aquí NO se toca;
  · checklist: título, pestañas, cabeceras de fase y las tareas de `datos_checklist.py` con el pie «Tareas
    completadas:» debajo; `Instrucciones!A3` con el nombre del producto; DV, formato condicional y contadores los rehace
    el motor.
Corre en el VPS (openpyxl). Después: `main.py --producto plan-negocio-sushi-bar --dry-run --origen <salida>` (ciclo.sh).

    /root/venv-guias/bin/python semilla_sus.py --salida /root/tmp-sus-f2/motor/origen

Sesión Claude Code, 31-oct-2026.
"""
import argparse
import copy
import importlib
import os
import shutil
import sys

import openpyxl

import datos_checklist as DC

AQUI = os.path.dirname(os.path.abspath(__file__))
REPO = os.path.abspath(os.path.join(AQUI, '..', '..', '..'))
TAQ = os.path.join(REPO, 'astro-site', 'public', 'dl', 'plan-negocio-taqueria')
PLAN_TAQ, CHECK_TAQ = 'plan-financiero-taqueria.xlsx', 'checklist-apertura-taqueria.xlsx'
PLAN, CHECK = 'plan-financiero-sushi-bar.xlsx', 'checklist-apertura-sushi-bar.xlsx'
PID = 'plan-negocio-sushi-bar'
NOMBRE = 'Plan de Negocio: Sushi Bar'

TITULOS_PLAN = {
    ('Inversión Inicial', 'A1'): 'PLAN DE NEGOCIO: SUSHI BAR',
    ('Inversión Inicial', 'A2'): 'Inversión inicial estimada — sushi bar con barra a la vista, mesas, para llevar y a '
                                 'domicilio',
    ('Inversión Inicial', 'A3'): 'Ajusta las cifras a tu proyecto: son las de un caso tipo en una ciudad media, sin '
                                 'traspaso. Pon tus presupuestos reales en las celdas de importe.',
    ('PyG 3 Años', 'A2'): 'Escenario realista — sushi bar con barra a la vista, mesas, para llevar y a domicilio',
    ('Punto Equilibrio', 'A2'): 'Punto de equilibrio mensual y diario — sushi bar con barra a la vista, mesas, para '
                                'llevar y a domicilio',
}

#: Filas 7-24 de 'Inversión Inicial': (rótulo, importe, nota, ¿lleva IVA?). R §3: 18 partidas, 98.500 € (obra 48.500 ·
#: maquinaria 39.000 · no amortizable 11.000) + imprevistos del 10 % sobre la obra (fila 27, por fórmula) = 103.350 €.
#: Sin palabras de recuento en el rótulo (motor.RX_RECUENTO), sin paréntesis con número, sin cifras de plazas. ⚠️
#: «campana» a secas la acentúa `motor.TILDES` a «campaña»: siempre «campana extractora»; y nunca el verbo «publica».
INVERSION = [
    ('Proyecto técnico de actividad y de extracción, y dirección de obra', 3500,
     'Proyecto de la actividad y del conducto de extracción, memoria para la licencia o la declaración responsable y '
     'certificado final de obra', 'Sí'),
    ('Obra del local con barra de sushi a la vista, rótulo y fachada', 28000,
     'Barra de sushi de obra a la vista, cocina alicatada, cuarto frío, suelo antideslizante, aseos (uno adaptado), '
     'falsos techos y aislamiento acústico, más el rótulo y la fachada. Estimación del kit: pide dos presupuestos', 'Sí'),
    ('Conducto de extracción exclusivo hasta cubierta', 5000,
     'Conducto exclusivo hasta la cubierta para la freidora de tempura y la plancha de las gyozas: Madrid pide superar '
     'en 1 m lo que haya en 15 m (Ordenanza 4/2021). Ningún fabricante da precio por metro: pide oferta con tu '
     'proyecto', 'Sí'),
    ('Instalación eléctrica, cuadro y boletín', 4500,
     'La cocina es eléctrica (arroceras, inducción, freidora y plancha), más el frío, la extracción y la '
     'climatización. Pregunta la potencia disponible ANTES de firmar: una acometida nueva tarda', 'Sí'),
    ('Fontanería, separador de grasas y climatización', 7500,
     'Agua fría y caliente en la cocina y en la barra, desagües con separador de grasas y climatización de la sala',
     'Sí'),
    ('Vitrina refrigerada de sushi para la barra', 2500,
     'La vitrina de la neta (el pescado cortado que va sobre el arroz) a la vista del cliente. Una de 6 cubetas GN 1/3 '
     'ronda los 999 € sin IVA y la de barra de gama alta pasa de 4.500 €: el caso cuenta una media', 'Sí'),
    ('Arroceras industriales, mantenedor de arroz y hangiri', 1500,
     'Dos arroceras (la de 8 litros ronda los 484 € sin IVA), un mantenedor y el hangiri, la tina de madera donde se '
     'aliña el shari. Estimación del kit', 'Sí'),
    ('Armario congelador de −20 °C con registrador y arcón', 2000,
     'El congelador del anisakis tiene que llevar el pescado a −20 °C o menos en la totalidad del producto durante 24 '
     'horas (RD 1021/2022, art. 8.1) y registrar la temperatura; un congelador que sólo marca −18 °C no sirve para '
     'eso. El arcón, para el resto del congelado. Estimación del kit: pide la ficha técnica', 'Sí'),
    ('Cocina caliente: inducción, freidora de tempura y plancha', 3500,
     'Placas de inducción para el miso y los caldos, freidora para la tempura y plancha para las gyozas. Estimación '
     'del kit', 'Sí'),
    ('Campana extractora con motor y filtros', 1500,
     'Campana extractora con motor y filtros de grasa sobre la freidora, la plancha y la inducción (una de 1,5-2 m '
     'ronda los 815 € sin IVA)', 'Sí'),
    ('Frío: armario de pescado, armarios de cocina, mesa refrigerada y nevera de bebida y sake', 7500,
     'Armario de pescado a 0-2 °C (el pescado fresco va a 4 °C o menos: RD 1021/2022, art. 4.1), armarios '
     'refrigerados de cocina, mesa refrigerada de trabajo para la barra y nevera de bebida y sake', 'Sí'),
    ('Lavavajillas de capota, fregaderos, lavamanos, mesas inox y estanterías', 5500,
     'Lavavajillas, fregaderos, lavamanos de accionamiento no manual, mesas de acero inoxidable y estanterías', 'Sí'),
    ('Utillaje: cuchillos japoneses, tablas por colores, sondas, básculas, pH-metro y etiquetadora', 3500,
     'Cuchillos yanagiba y deba, tablas por colores, termómetros sonda, básculas, el pH-metro para medir el shari de '
     'cada tanda, la etiquetadora de las bandejas cerradas (RD 126/2015, art. 5), vajilla y la cartelería de '
     'alérgenos', 'Sí'),
    ('Mobiliario de sala y barra', 9000,
     'Mesas y sillas para 28 plazas y 8 taburetes en la barra de sushi', 'Sí'),
    ('TPV con Veri*factu, datáfono, impresora de comandas y tablets', 2500,
     'Software de facturación conforme a Veri*factu, datáfono, impresora de cocina y las tablets de las plataformas',
     'Sí'),
    ('Constitución de la SL, licencia, tasas, registro sanitario y APPCC', 3000,
     'SL unipersonal, licencia o declaración responsable, tasas municipales, registro sanitario autonómico y '
     'redacción del plan APPCC con el anisakis y el shari. Las tasas no llevan IVA; los honorarios sí: cambia la '
     'columna de al lado si tu presupuesto es sobre todo de honorarios', 'No'),
    ('Marca, web, fotos y lanzamiento', 4500,
     'Identidad, web con la carta y los alérgenos, sesión de fotos y campaña de apertura', 'Sí'),
    ('Stock inicial: pescado, arroz, algas, salsas, cajas de pedidos y bebida', 3500,
     'Pescado, arroz para sushi, nori, salsas, el primer pedido de cajas para llevar y a domicilio y la bebida, el sake '
     'incluido (pídelo con su plazo: depende de un importador)', 'Sí'),
]
assert len(INVERSION) == 18 and sum(x[1] for x in INVERSION) == 98500, sum(x[1] for x in INVERSION)

#: Filas 26-34 de 'PyG 3 Años' (costes fijos propios): 9 filas, una por fila del molde (no se suprimen ni se añaden:
#: desplazaría el P&L), 13.850 €/año (R §3). Sin palabras de recuento en el rótulo. Las claves de `contenido….FIJOS`
#: son estos rótulos normalizados.
FIJOS = {
    26: ('Gestoría laboral, fiscal y contable', 2400),
    27: ('Mantenimiento del frío, el congelador, la campana extractora y el pH-metro', 1800),
    28: ('Marketing y publicidad', 4000),
    29: ('Software, TPV, teléfono e internet', 1800),
    30: ('Limpieza, lavandería, EPI y análisis de autocontrol', 1500),
    31: ('Gestión de residuos y del aceite usado de la freidora', 600),
    32: ('Desinsectación, desratización y desinfección (DDD)', 450),
    33: ('Derechos de autor por música ambiental (SGAE y AGEDI-AIE)', 500),
    34: ('Prevención de riesgos laborales, vigilancia de la salud y formación', 800),
}
assert sum(x[1] for x in FIJOS.values()) == 13850


def _contenido():
    sys.path.insert(0, os.path.join(os.path.dirname(AQUI), 'planes-v2_0'))
    import grupo_a
    return grupo_a, importlib.import_module('contenido_plan_negocio_sushi_bar.a')


def _supuestos(wb):
    """Los valores de `0. Supuestos` del caso sushi bar: `grupo_a` CONSERVA lo que el libro ya trae en las celdas
    verdes y sólo usa el valor del contenido en una celda vacía; sin este paso el caso seguiría siendo el de la
    taquería (77 × 15,50 € × 300, 55.000 + 118.000 €)."""
    grupo_a, contenido = _contenido()
    coords = dict((t[0], t[1]) for t in grupo_a.SUPUESTOS_BASE)
    ws = wb['0. Supuestos']
    escritas = 0
    for clave, t in contenido.SUPUESTOS.items():
        if t[2] is None:
            continue
        assert clave in coords, clave
        ws[coords[clave]] = t[2]
        escritas += 1
    # F1-camino §7 F2 (2): las pagas con el Mes 1 en enero ya vienen de la Taquería (7/12)
    assert coords['mes_paga_1'] == 'B60' and coords['mes_paga_2'] == 'B61', (coords['mes_paga_1'], coords['mes_paga_2'])
    assert (ws['B60'].value, ws['B61'].value) == (7, 12), (ws['B60'].value, ws['B61'].value)
    return escritas


def _docprops(wb):
    """El motor no reescribe `description` ni `keywords` y heredaba los de la taquería."""
    wb.properties.description = 'aichef.pro/' + PID
    wb.properties.keywords = 'plan de negocio sushi bar, restaurante de sushi, sushiman, AI Chef Pro'


def plan(salida):
    destino = os.path.join(salida, PLAN)
    shutil.copy2(os.path.join(TAQ, PLAN_TAQ), destino)
    wb = openpyxl.load_workbook(destino)
    print('  supuestos escritos:', _supuestos(wb))
    _docprops(wb)
    for (hoja, coord), texto in TITULOS_PLAN.items():
        wb[hoja][coord] = texto
    ws = wb['Inversión Inicial']
    for i, (rot, imp, nota, iva) in enumerate(INVERSION):
        r = 7 + i
        assert not str(ws.cell(r, 1).value).lower().startswith(('fianza', 'alquiler', 'imprevistos')), r
        ws.cell(r, 1).value, ws.cell(r, 2).value, ws.cell(r, 4).value, ws.cell(r, 5).value = rot, imp, nota, iva
    assert str(ws['A25'].value).lower().startswith('fianza'), ws['A25'].value
    assert str(ws['A26'].value).lower().startswith('alquiler'), ws['A26'].value
    assert str(ws['A27'].value).lower().startswith('imprevistos'), ws['A27'].value
    ws = wb['PyG 3 Años']
    assert str(ws['A25'].value).startswith('Seguros') and str(ws['A35'].value).startswith('Amortización'), (
        ws['A25'].value, ws['A35'].value)
    for r, (rot, imp) in FIJOS.items():
        ws.cell(r, 1).value, ws.cell(r, 2).value = rot, imp
        ws.cell(r, 6).value = None          # la nota la pone el contenido (FIJOS)
    # `PyG!G13` (IVA soportado de la compra de comida): en la Taquería ya es VALOR (0,085) y el motor lo reescribe con
    # `IVA_COMPRAS` (0,087). No se toca aquí.
    assert str(ws['A13'].value).startswith('Coste de mercancía — comida'), ws['A13'].value
    assert not str(ws['G13'].value).startswith('='), ws['G13'].value
    # La rotación de la Taquería se queda ENCENDIDA: sus fórmulas tienen que seguir en el origen (F1-camino §0.1)
    assert str(wb['0. Supuestos']['B50'].value).startswith('='), wb['0. Supuestos']['B50'].value
    wb.save(destino)
    return destino


def _copiar_estilo(origen, destino):
    destino.font, destino.border = copy.copy(origen.font), copy.copy(origen.border)
    destino.fill, destino.alignment = copy.copy(origen.fill), copy.copy(origen.alignment)
    destino.number_format, destino.protection = origen.number_format, copy.copy(origen.protection)


def checklist(salida):
    destino = os.path.join(salida, CHECK)
    shutil.copy2(os.path.join(TAQ, CHECK_TAQ), destino)
    wb = openpyxl.load_workbook(destino)
    fases = [ws for ws in wb.worksheets if ws.title[:2] in ('F1', 'F2', 'F3', 'F4', 'F5', 'F6')]
    assert len(fases) == 6
    for ws, (titulo, cabecera, tareas) in zip(fases, DC.FASES):
        ws.title = titulo
        ws['A1'], ws['A2'] = DC.TITULO, cabecera
        assert ws['B4'].value and 'TRÁMITE' in ws['B4'].value.upper(), ws['B4'].value
        modelo = [copy.copy(ws.cell(5, c)) for c in range(1, 6)]
        alto = ws.row_dimensions[5].height
        pie = [copy.copy(ws.cell(r, c)) for r in range(5, ws.max_row + 1) for c in range(1, 6)
               if isinstance(ws.cell(r, c).value, str) and 'completad' in ws.cell(r, c).value.lower()]
        assert pie, ws.title
        fila_pie = pie[0].row
        estilo_pie = [copy.copy(ws.cell(fila_pie, c)) for c in range(1, 6)]
        # vaciar la zona de tareas y el pie (valores, DV y formato condicional los rehace el motor)
        for r in range(5, ws.max_row + 1):
            for c in range(1, 6):
                ws.cell(r, c).value = None
            ws.row_dimensions[r].height = None
        ws.data_validations.dataValidation = []
        ws.conditional_formatting = type(ws.conditional_formatting)()
        for i, (tarea, resp, plazo, nota) in enumerate(tareas):
            r = 5 + i
            for c, v in enumerate(('☐', tarea, resp, plazo, nota), start=1):
                cel = ws.cell(r, c)
                _copiar_estilo(modelo[c - 1], cel)
                cel.value = v
            ws.row_dimensions[r].height = alto
        r = 5 + len(tareas) + 1
        for c in range(1, 6):
            _copiar_estilo(estilo_pie[c - 1], ws.cell(r, c))
        ws.cell(r, 2).value = 'Tareas completadas:'
        ws.cell(r, 3).value = '=COUNTIF(A5:A%d,"✓")' % (4 + len(tareas))
        ws.cell(r, 4).value = 'de'
        ws.cell(r, 5).value = '=COUNTIF(B5:B%d,"?*")-COUNTIF(A5:A%d,"N/A")' % (4 + len(tareas), 4 + len(tareas))
    # `motor.asegurar_instrucciones` no reescribe la hoja si ya existe: la de la Taquería nombraba su producto
    ins = wb['Instrucciones']
    assert 'Taquería' in str(ins['A3'].value), ins['A3'].value
    ins['A3'] = 'Este fichero forma parte del «%s» de AI Chef Pro.' % NOMBRE
    _docprops(wb)
    wb.save(destino)
    return destino


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--salida', required=True)
    a = ap.parse_args()
    if os.path.isdir(a.salida):
        shutil.rmtree(a.salida)
    os.makedirs(a.salida)
    print('plan      →', plan(a.salida))
    print('checklist →', checklist(a.salida), '(%d tareas: %s)' % (
        sum(len(t) for _, _, t in DC.FASES), '/'.join(str(len(t)) for _, _, t in DC.FASES)))


if __name__ == '__main__':
    main()
