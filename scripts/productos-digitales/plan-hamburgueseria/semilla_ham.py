#!/usr/bin/env python3
"""
semilla_ham.py — Plan de Negocio: Hamburguesería Smash (ES 1.0; F1-camino §2-§3): prepara el ORIGEN que
`planes-v2_0/main.py` convierte en el plan financiero y el checklist de la hamburguesería. Copia adaptada de
`plan-sushi/semilla_sus.py` (la misma rejilla: la rotación se queda ENCENDIDA, F1-camino §0.1).

Copia los dos xlsx del Sushi Bar ES 1.0 publicados (`dl/plan-negocio-sushi-bar/`, SOLO plan y checklist, SOLO
LECTURA; molde A-β + checklist C2; su ciclo reproduce el libro LIVE con 0 diferencias, F1-camino §1) con los nombres de
fichero de la hamburguesería y escribe lo que el módulo de contenido no alcanza:
  · plan: valores de `0. Supuestos` (grupo_a conserva lo verde que trae el libro; `B60`/`B61` ya son 7/12 en el
    Sushi Bar: assert, sin escribir), títulos de hoja (Inversión A1/A2/A3, PyG A2, Punto Equilibrio A2), las 18
    partidas de inversión (filas 7-24: rótulo, importe, nota, ¿IVA?), los 9 costes fijos propios del P&L (filas
    26-34: rótulo e importe; la nota la pone el contenido) y los docProps description/keywords. `PyG!G13` ya es VALOR
    en el Sushi Bar y lo rehace el motor con `IVA_COMPRAS` (0,090): aquí NO se toca;
  · checklist: título, pestañas, cabeceras de fase y las tareas de `datos_checklist.py` con el pie «Tareas
    completadas:» debajo; `Instrucciones!A3` con el nombre del producto; DV, formato condicional y contadores los rehace
    el motor.
Corre en el VPS (openpyxl). Después: `main.py --producto plan-negocio-hamburgueseria --dry-run --origen <salida>`
(ciclo.sh).

    /root/venv-guias/bin/python semilla_ham.py --salida /root/tmp-ham-f2/motor/origen

Sesión Claude Code, 31-oct-2026.
"""
import argparse
import copy
import importlib
import os
import shutil
import sys

import openpyxl

import datos_checklist as DC

AQUI = os.path.dirname(os.path.abspath(__file__))
REPO = os.path.abspath(os.path.join(AQUI, '..', '..', '..'))
SUS = os.path.join(REPO, 'astro-site', 'public', 'dl', 'plan-negocio-sushi-bar')
PLAN_SUS, CHECK_SUS = 'plan-financiero-sushi-bar.xlsx', 'checklist-apertura-sushi-bar.xlsx'
PLAN, CHECK = 'plan-financiero-hamburgueseria.xlsx', 'checklist-apertura-hamburgueseria.xlsx'
PID = 'plan-negocio-hamburgueseria'
NOMBRE = 'Plan de Negocio: Hamburguesería Smash'
KEYWORDS = ('plan de negocio hamburguesería, cómo montar una hamburguesería, smash burger, AI Chef Pro')

TITULOS_PLAN = {
    ('Inversión Inicial', 'A1'): 'PLAN DE NEGOCIO: HAMBURGUESERÍA SMASH',
    ('Inversión Inicial', 'A2'): 'Inversión inicial estimada — hamburguesería de smash burgers con sala, para llevar y '
                                 'a domicilio',
    ('Inversión Inicial', 'A3'): 'Ajusta las cifras a tu proyecto: son las de un caso tipo en una ciudad media, sin '
                                 'traspaso. Pon tus presupuestos reales en las celdas de importe.',
    ('PyG 3 Años', 'A2'): 'Escenario realista — hamburguesería de smash burgers con sala, para llevar y a domicilio',
    ('Punto Equilibrio', 'A2'): 'Punto de equilibrio mensual y diario — hamburguesería de smash burgers con sala, para '
                                'llevar y a domicilio',
}

#: Filas 7-24 de 'Inversión Inicial': (rótulo, importe, nota, ¿lleva IVA?). R §3: 18 partidas, 99.500 € (obra 48.500 ·
#: maquinaria 40.000 · no amortizable 11.000) + imprevistos del 10 % sobre la obra (fila 27, por fórmula) = 104.350 €.
#: Sin palabras de recuento en el rótulo (motor.RX_RECUENTO), sin paréntesis con número, sin cifras de plazas (la fila
#: 14 no dice «36 plazas»; el «≤ 2 °C» de la fila 10 va a la nota). ⚠️ «campana» a secas la acentúa `motor.TILDES` a
#: «campaña»: siempre «campana extractora»; y nunca el verbo «publica».
INVERSION = [
    ('Proyecto técnico de actividad y de extracción, y dirección de obra', 3500,
     'Proyecto de la actividad y del conducto de extracción, memoria para la licencia o la declaración responsable y '
     'certificado final de obra', 'Sí'),
    ('Obra del local con mostrador de pedidos y cocina a la vista, rótulo y fachada', 26000,
     'Mostrador de pedidos con la cocina a la vista, cocina alicatada, suelo antideslizante, aseos (uno adaptado), '
     'falsos techos y aislamiento acústico, más el rótulo y la fachada. Estimación del kit: pide dos presupuestos', 'Sí'),
    ('Conducto de extracción exclusivo hasta cubierta', 6000,
     'Conducto exclusivo hasta la cubierta para la plancha y las freidoras, que echan mucha grasa: Madrid pide superar '
     'en 1 m lo que haya en 15 m (Ordenanza 4/2021). Ningún fabricante da precio por metro: pide oferta con tu '
     'proyecto', 'Sí'),
    ('Instalación eléctrica, cuadro y boletín', 4500,
     'Las freidoras eléctricas suelen pedir corriente trifásica, más el frío, la extracción y la climatización. '
     'Pregunta la potencia disponible ANTES de firmar: una acometida nueva tarda', 'Sí'),
    ('Fontanería, separador de grasas, climatización e instalación de gas', 8500,
     'Agua fría y caliente en la cocina, desagües con separador de grasas, climatización de la sala y la instalación '
     'de gas de la plancha por una empresa instaladora habilitada, con su certificado (RD 919/2006)', 'Sí'),
    ('Plancha fry-top de acero a gas y tostador de pan', 4000,
     'La plancha de acero (fry-top) donde se aplasta cada bola: una de 60 a 80 cm va de unos 771 a 1.399 € sin IVA según la '
     'gama; el caso cuenta una de 100 cm, más el tostador del pan. Estimación del kit', 'Sí'),
    ('Dos freidoras eléctricas de doble cuba y mantenedor de patatas', 3500,
     'Dos freidoras para no mezclar las patatas con los rebozados (alérgenos y contacto cruzado) y el mantenedor de '
     'patatas. Las de gama ligera bajan de 300 € sin IVA, pero aguantan poco en un servicio fuerte: pide la ficha '
     'técnica', 'Sí'),
    ('Campana extractora con motor y filtros', 2200,
     'Campana extractora con motor y filtros de grasa sobre la plancha y las freidoras (una de 1,5 a 2 m va de unos '
     '543 a 813 € sin IVA; el caso cuenta una de 3 m)', 'Sí'),
    ('Mesa refrigerada de montaje y bajo plancha refrigerado para las bolas', 3800,
     'La línea de montaje con los ingredientes en frío y el bajo plancha con cajones para tener las bolas a mano y en '
     'frío durante el servicio (una mesa refrigerada va de unos 1.177 a 1.948 € sin IVA)', 'Sí'),
    ('Frío: armario de carne, armarios de cocina, congelador de patatas y nevera de bebidas', 5000,
     'Armario de carne a 2 °C o menos (la carne picada va a esa temperatura y los preparados de carne, a 4 °C o menos: '
     'RD 1021/2022, art. 4.1), armarios refrigerados de cocina, congelador de patatas y nevera de bebidas', 'Sí'),
    ('Barra: batidora de batidos, arcón de helado, tirador de cerveza y botellero', 2500,
     'La batidora y el arcón de helado de los batidos, el tirador de cerveza (si no te lo cede la cervecera) y el '
     'botellero', 'Sí'),
    ('Lavavajillas de capota, fregaderos, lavamanos, mesas inox y estanterías', 5500,
     'Lavavajillas, fregaderos, lavamanos de accionamiento no manual, mesas de acero inoxidable y estanterías', 'Sí'),
    ('Utillaje: smashers, espátulas, básculas, sondas, medidor de compuestos polares y etiquetadora', 2000,
     'Smashers (las prensas para aplastar la bola), espátulas y raspadores, básculas de porcionado, termómetros sonda, '
     'el medidor de compuestos polares del aceite (de 399 a 499 € sin IVA; tu APPCC puede elegir tiras), cubetas GN, '
     'dispensadores de salsa y la etiquetadora', 'Sí'),
    ('Mobiliario de sala, barra de ventana y mostrador', 8000,
     'Mesas y sillas para 30 plazas, 6 taburetes en la barra de ventana y el mueble del mostrador', 'Sí'),
    ('TPV con Veri*factu, pantalla de turnos, datáfono, impresora de comandas e integrador de plataformas', 3500,
     'Software de facturación conforme a Veri*factu, la pantalla que llama los turnos, datáfono, impresora de cocina y '
     'el integrador que junta los pedidos de las plataformas', 'Sí'),
    ('Constitución de la SL, licencia, tasas, registro sanitario y APPCC', 3000,
     'SL unipersonal, licencia o declaración responsable, tasas municipales, registro sanitario autonómico y '
     'redacción del plan APPCC con la carne picada, la cocción y la freidora. Las tasas no llevan IVA; los honorarios '
     'sí: cambia la columna de al lado si tu presupuesto es sobre todo de honorarios', 'No'),
    ('Marca, web, fotos y vídeo de lanzamiento', 5000,
     'Identidad, web con la carta, los alérgenos y los pedidos para recoger, sesión de fotos y el vídeo de la '
     'apertura (la smash se vende en vídeo)', 'Sí'),
    ('Stock inicial: carne, pan, patatas, aceite, salsas, envases y bebida', 3000,
     'Carne, pan, patatas, aceite, salsas, el primer pedido de cajas y bolsas para llevar y a domicilio y la bebida',
     'Sí'),
]
assert len(INVERSION) == 18 and sum(x[1] for x in INVERSION) == 99500, sum(x[1] for x in INVERSION)

#: Filas 26-34 de 'PyG 3 Años' (costes fijos propios): 9 filas, una por fila del molde (no se suprimen ni se añaden:
#: desplazaría el P&L), 14.400 €/año (R §3). Sin palabras de recuento en el rótulo. Las claves de `contenido….FIJOS`
#: son estos rótulos normalizados.
FIJOS = {
    26: ('Gestoría laboral, fiscal y contable', 2400),
    27: ('Mantenimiento de la campana extractora y el conducto, la plancha y el gas, las freidoras y el frío', 2000),
    28: ('Marketing y publicidad', 4500),
    29: ('Software, TPV, integrador de plataformas, teléfono e internet', 1800),
    30: ('Limpieza, lavandería, EPI y análisis de autocontrol', 1500),
    31: ('Gestión de residuos y del aceite usado de las freidoras', 600),
    32: ('Desinsectación, desratización y desinfección (DDD)', 450),
    33: ('Derechos de autor por música ambiental (SGAE y AGEDI-AIE)', 500),
    34: ('Prevención de riesgos laborales, vigilancia de la salud y formación', 650),
}
assert sum(x[1] for x in FIJOS.values()) == 14400


def _contenido():
    sys.path.insert(0, os.path.join(os.path.dirname(AQUI), 'planes-v2_0'))
    import grupo_a
    return grupo_a, importlib.import_module('contenido_plan_negocio_hamburgueseria.a')


def _supuestos(wb):
    """Los valores de `0. Supuestos` del caso hamburguesería: `grupo_a` CONSERVA lo que el libro ya trae en las celdas
    verdes y sólo usa el valor del contenido en una celda vacía; sin este paso el caso seguiría siendo el del sushi bar
    (79 × 20,00 € × 300, 55.000 + 131.000 €)."""
    grupo_a, contenido = _contenido()
    coords = dict((t[0], t[1]) for t in grupo_a.SUPUESTOS_BASE)
    ws = wb['0. Supuestos']
    escritas = 0
    for clave, t in contenido.SUPUESTOS.items():
        if t[2] is None:
            continue
        assert clave in coords, clave
        ws[coords[clave]] = t[2]
        escritas += 1
    # F1-camino §7 F2 (2): las pagas con el Mes 1 en enero ya vienen del Sushi Bar (7/12)
    assert coords['mes_paga_1'] == 'B60' and coords['mes_paga_2'] == 'B61', (coords['mes_paga_1'], coords['mes_paga_2'])
    assert (ws['B60'].value, ws['B61'].value) == (7, 12), (ws['B60'].value, ws['B61'].value)
    return escritas


def _docprops(wb):
    """El motor no reescribe `description` ni `keywords` y heredaba los del sushi bar."""
    wb.properties.description = 'aichef.pro/' + PID
    wb.properties.keywords = KEYWORDS


def plan(salida):
    destino = os.path.join(salida, PLAN)
    shutil.copy2(os.path.join(SUS, PLAN_SUS), destino)
    wb = openpyxl.load_workbook(destino)
    print('  supuestos escritos:', _supuestos(wb))
    _docprops(wb)
    for (hoja, coord), texto in TITULOS_PLAN.items():
        wb[hoja][coord] = texto
    ws = wb['Inversión Inicial']
    for i, (rot, imp, nota, iva) in enumerate(INVERSION):
        r = 7 + i
        assert not str(ws.cell(r, 1).value).lower().startswith(('fianza', 'alquiler', 'imprevistos')), r
        ws.cell(r, 1).value, ws.cell(r, 2).value, ws.cell(r, 4).value, ws.cell(r, 5).value = rot, imp, nota, iva
    assert str(ws['A25'].value).lower().startswith('fianza'), ws['A25'].value
    assert str(ws['A26'].value).lower().startswith('alquiler'), ws['A26'].value
    assert str(ws['A27'].value).lower().startswith('imprevistos'), ws['A27'].value
    ws = wb['PyG 3 Años']
    assert str(ws['A25'].value).startswith('Seguros') and str(ws['A35'].value).startswith('Amortización'), (
        ws['A25'].value, ws['A35'].value)
    for r, (rot, imp) in FIJOS.items():
        ws.cell(r, 1).value, ws.cell(r, 2).value = rot, imp
        ws.cell(r, 6).value = None          # la nota la pone el contenido (FIJOS)
    # `PyG!G13` (IVA soportado de la compra de comida): en el Sushi Bar ya es VALOR (0,087) y el motor lo reescribe con
    # `IVA_COMPRAS` (0,090). No se toca aquí.
    assert str(ws['A13'].value).startswith('Coste de mercancía — comida'), ws['A13'].value
    assert not str(ws['G13'].value).startswith('='), ws['G13'].value
    # La rotación del Sushi Bar se queda ENCENDIDA: sus fórmulas tienen que seguir en el origen (F1-camino §0.1)
    assert str(wb['0. Supuestos']['B50'].value).startswith('='), wb['0. Supuestos']['B50'].value
    wb.save(destino)
    return destino


def _copiar_estilo(origen, destino):
    destino.font, destino.border = copy.copy(origen.font), copy.copy(origen.border)
    destino.fill, destino.alignment = copy.copy(origen.fill), copy.copy(origen.alignment)
    destino.number_format, destino.protection = origen.number_format, copy.copy(origen.protection)


def checklist(salida):
    destino = os.path.join(salida, CHECK)
    shutil.copy2(os.path.join(SUS, CHECK_SUS), destino)
    wb = openpyxl.load_workbook(destino)
    fases = [ws for ws in wb.worksheets if ws.title[:2] in ('F1', 'F2', 'F3', 'F4', 'F5', 'F6')]
    assert len(fases) == 6
    for ws, (titulo, cabecera, tareas) in zip(fases, DC.FASES):
        ws.title = titulo
        ws['A1'], ws['A2'] = DC.TITULO, cabecera
        assert ws['B4'].value and 'TRÁMITE' in ws['B4'].value.upper(), ws['B4'].value
        modelo = [copy.copy(ws.cell(5, c)) for c in range(1, 6)]
        alto = ws.row_dimensions[5].height
        pie = [copy.copy(ws.cell(r, c)) for r in range(5, ws.max_row + 1) for c in range(1, 6)
               if isinstance(ws.cell(r, c).value, str) and 'completad' in ws.cell(r, c).value.lower()]
        assert pie, ws.title
        fila_pie = pie[0].row
        estilo_pie = [copy.copy(ws.cell(fila_pie, c)) for c in range(1, 6)]
        # vaciar la zona de tareas y el pie (valores, DV y formato condicional los rehace el motor)
        for r in range(5, ws.max_row + 1):
            for c in range(1, 6):
                ws.cell(r, c).value = None
            ws.row_dimensions[r].height = None
        ws.data_validations.dataValidation = []
        ws.conditional_formatting = type(ws.conditional_formatting)()
        for i, (tarea, resp, plazo, nota) in enumerate(tareas):
            r = 5 + i
            for c, v in enumerate(('☐', tarea, resp, plazo, nota), start=1):
                cel = ws.cell(r, c)
                _copiar_estilo(modelo[c - 1], cel)
                cel.value = v
            ws.row_dimensions[r].height = alto
        r = 5 + len(tareas) + 1
        for c in range(1, 6):
            _copiar_estilo(estilo_pie[c - 1], ws.cell(r, c))
        ws.cell(r, 2).value = 'Tareas completadas:'
        ws.cell(r, 3).value = '=COUNTIF(A5:A%d,"✓")' % (4 + len(tareas))
        ws.cell(r, 4).value = 'de'
        ws.cell(r, 5).value = '=COUNTIF(B5:B%d,"?*")-COUNTIF(A5:A%d,"N/A")' % (4 + len(tareas), 4 + len(tareas))
    # `motor.asegurar_instrucciones` no reescribe la hoja si ya existe: la del Sushi Bar nombraba su producto
    ins = wb['Instrucciones']
    assert 'Sushi Bar' in str(ins['A3'].value), ins['A3'].value
    ins['A3'] = 'Este fichero forma parte del «%s» de AI Chef Pro.' % NOMBRE
    _docprops(wb)
    wb.save(destino)
    return destino


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--salida', required=True)
    a = ap.parse_args()
    if os.path.isdir(a.salida):
        shutil.rmtree(a.salida)
    os.makedirs(a.salida)
    print('plan      →', plan(a.salida))
    print('checklist →', checklist(a.salida), '(%d tareas: %s)' % (
        sum(len(t) for _, _, t in DC.FASES), '/'.join(str(len(t)) for _, _, t in DC.FASES)))


if __name__ == '__main__':
    main()
