#!/usr/bin/env python3
"""barrido_es.py — Plan de Negocio: Hamburguesería Smash ES 1.0 (F1-camino §4 y §7; copia de `plan-sushi/barrido_es.py`):
barre los xlsx/docx que se le pasen buscando los prohibidos (España 2026, validado, freelance, carné de manipulador, tarifa
plana, CNAE, RGSEAA para el minorista, licencia de apertura, RD 895/2013, RD 393/2007, neto «15-25 %», versión
anterior…), los restos de los moldes (Food Truck, Tapas Bar, Heladería, Pizzería, Taquería, Pollería y el Sushi Bar —la
semilla—), los errores legales de la carne (75 °C, 155 °F o 68 °C) y las frases del motor sin traducir: la lista es
`tandas_docx.PROHIBIDOS`, y antes se borra `tandas_docx.PERMITIR` («Plan de Negocio: Food Truck», la variante D).
Además, como ERROR: «Campaña» por «campana» (`RX_CAMPANA`), «pública» por el verbo «publica» (`RX_PUBLICA`) y las citas
legales mal de `tandas_docx.art30_mal` («art. 30» sin «RD 1086/2020» o atribuido al RD 1021/2022, «7 días» junto a
congelación, picada, bola o −18 °C, una temperatura de cocción «legal» u «obligatoria» sin negación y «de un día para
otro» sin negación: F1-camino §7 F2 (5)). Y como AVISO, uno a uno, para revisar a mano (no cambian el exit): −18 °C,
carne madurada, «al punto» o «poco hecha», «sin gluten», «auténtic», «gourmet», batido junto al 21 %, Día de la
Hamburguesa, «digital obligatorio», franquicia o Goiko, 180-190 °C, «pan común», «recongelar» y «medidor» junto a
«obligatori» (`tandas_docx.AVISOS`). Lista celda/párrafo y texto; en las fórmulas sólo mira los literales entre
comillas. Corre en el VPS. Uso: barrido_es.py <fichero>… · exit 1 si hay coincidencias de error.
Sesión Claude Code, 31-oct-2026."""
import os
import re
import sys

import docx
import openpyxl

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import tandas_docx as T                                                            # noqa: E402

#: «Campaña» por «campana» (motor.TILDES acentúa «campana» suelta; Pizzería F2-NOTAS §2.8): la única «campaña» legítima es
#: la de apertura, lanzamiento, marketing, publicidad o Navidad.
RX_CAMPANA = r'[Cc]ampa[ñÑ]a(?! de (apertura|lanzamiento|marketing|publicidad|Navidad))'
#: «pública» por «publica» (la misma trampa con el verbo; revisión final de la Pizzería, m1).
RX_PUBLICA = r'\b(nadie|se|le|les|te|nos|quien) públic[ao]n?\b|\bpúblic[ao] (precios?|tarifas?)\b'
RX = re.compile('|'.join(T.PROHIBIDOS + [RX_CAMPANA, RX_PUBLICA]), re.I)
n = avisos = 0
for p in [x for x in sys.argv[1:] if not x.startswith('--')]:
    if p.endswith('.xlsx'):
        wb = openpyxl.load_workbook(p)
        it = [(ws.title + '!' + c.coordinate, c.value) for ws in wb.worksheets for row in ws.iter_rows() for c in row
              if isinstance(c.value, str)]
        it += [('pestaña', ws.title) for ws in wb.worksheets]
        pr = wb.properties
        it += [('docProps.' + k, getattr(pr, k)) for k in ('title', 'subject', 'description', 'keywords', 'creator')
               if isinstance(getattr(pr, k), str)]
    else:
        d = docx.Document(p)
        it = [('p%d' % i, x.text) for i, x in enumerate(d.paragraphs)]
        it += [('tabla%d' % i, c.text) for i, t in enumerate(d.tables) for row in t.rows for c in row.cells]
        cp = d.core_properties
        it += [('docProps.' + k, getattr(cp, k)) for k in ('title', 'subject', 'comments', 'keywords', 'author')
               if isinstance(getattr(cp, k), str)]
    nombre = p.split('/')[-1][:28]
    for loc, t in it:
        txt = ' '.join(re.findall(r'"([^"]*)"', t)) if t.startswith('=') else t
        m = RX.search(T.PERMITIR.sub('', txt))
        if m:
            n += 1
            tp = T.PERMITIR.sub('', txt)
            print('%s · %s · «%s» · …%s…' % (nombre, loc, m.group(0), tp[max(0, m.start() - 50):m.end() + 50]))
        for frag in T.art30_mal(txt):
            n += 1
            print('%s · %s · cita legal mal · …%s…' % (nombre, loc, frag))
        for etiqueta, rx in T.AVISOS:
            m = rx.search(txt)
            if m:
                avisos += 1
                print('AVISO %s · %s · %s · …%s…' % (nombre, loc, etiqueta, txt[max(0, m.start() - 60):m.end() + 60]))
print('barrido: %d coincidencias · %d avisos para revisar a mano' % (n, avisos))
sys.exit(1 if n else 0)
