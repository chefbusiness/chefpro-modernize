#!/usr/bin/env python3
"""
calculadora_ham_es.py — Plan de Negocio: Hamburguesería Smash ES 1.0: `calculadora-carne-patatas-escandallo-hamburgueseria.xlsx`.
Hojas `Carne` · `Patatas` · `Escandallo` · `Instrucciones`. Gemelo de código: `../plan-sushi/calculadora_sus_es.py` (helpers,
inject_cache, autotest con pycel y un `calcular()` en Python, protección sin contraseña, verdes desbloqueadas, A4, docProps
AI Chef Pro). Todos los valores por defecto son «estimación del kit, editable» (F1-research-hamburgueseria.md §2 y §8, con los tres
ajustes de F1-camino.md §4: queso de 20 g por bola, bolas por kilo en crudo y con la merma de cierre, y «bolas por comensal»).
  · Carne: 3 carnes + 1 libre (blend comprado ya picado, pieza para picar en el local, bolas porcionadas) → coste por kg útil =
    precio ÷ útil ÷ (1 − merma de cierre), coste por bola (de 1, 2 y 3), bolas por kilo, kg de carne al día y «picar en el
    local te compensa mientras limpiarlo cueste menos de X € por kg útil». Desplegable informativo de carne picada.
  · Patatas: ración de patata congelada y fresca (con y sin mano de obra), coste del aceite por ración, kg de patata al día y
    litros de aceite al mes, y «la fresca te compensa mientras pelar y cortar cueste menos de X € por kg».
  · Escandallo: 14 productos × (bolas, lonchas de queso, panes, otros ingredientes, raciones de patata, PVP, ventas a la semana) →
    coste, food cost con semáforo, PVP sugerido, margen por canal (sala · para llevar · plataforma), puesto por margen y food cost
    ponderado.
Valores que otros scripts leen de la CACHÉ del xlsx (dict `EXPUESTOS`, función `valores_expuestos(ruta)`).
Corre en el VPS (openpyxl + pycel):

    /root/venv-guias/bin/python calculadora_ham_es.py --salida <dir>      # genera + inject_cache + autotest
    /root/venv-guias/bin/python calculadora_ham_es.py --salida <dir> --autotest

Sesión Claude Code, 31-oct-2026.
"""
import argparse
import math
import os
import re
import sys
from collections import OrderedDict

import openpyxl
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill, Protection
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

AQUI = os.path.dirname(os.path.abspath(__file__))
PD = os.path.dirname(AQUI)
sys.path.insert(0, AQUI)
sys.path.insert(0, os.path.join(PD, 'planes-v2_0'))
import motor                                                                       # noqa: E402

# Versión del producto NUEVO: 1.0 · noviembre 2026 (mismo parche que sus_version.py / ham_version.py, sin tocar el motor común)
motor.VERSION = '1.0'
motor.VERSION_MES = 'noviembre 2026'
motor.VERSION_FMT = 'Versión ' + motor.VERSION + ' · ' + motor.VERSION_MES + ' · aichef.pro/{pid} · info@aichef.pro'

PID = 'plan-negocio-hamburgueseria'
FICHERO = 'calculadora-carne-patatas-escandallo-hamburgueseria.xlsx'
TITULO = 'Calculadora de la carne, las patatas y el escandallo'
HC, HP, HE, HI = 'Carne', 'Patatas', 'Escandallo', 'Instrucciones'
VERDE = PatternFill('solid', fgColor='E8F5E9')
GRIS = PatternFill('solid', fgColor='F2F2F2')
ROJO = PatternFill('solid', bgColor='F8CBAD', fgColor='F8CBAD')
AMBAR = PatternFill('solid', bgColor='FFE699', fgColor='FFE699')
OK = PatternFill('solid', bgColor='C6E0B4', fgColor='C6E0B4')
FUENTE = 'Calibri'
EUR, ENT, PCT, DEC, GR, EUR3 = '#,##0.00 €', '#,##0', '0.0%', '#,##0.0', '#,##0.0', '#,##0.000 €'
FC_OK, FC_VIG, FC_ALTO = 'En objetivo', 'Vigila el coste', 'Por encima del objetivo'
SEM_PICAR, SEM_BLEND = 'Picar en el local compensa', 'Compra el blend ya picado'
SEM_FRE, SEM_CONG = 'La fresca compensa', 'La congelada compensa'
EST = ' Estimación del kit, editable.'

# ---------------------------------------------------------------------------------------------------------------------
# Datos (F1-research §2 y §8). Importes sin IVA salvo el PVP de carta.
# ---------------------------------------------------------------------------------------------------------------------
NPE = 4
PC = [get_column_letter(2 + i) for i in range(NPE)]                                # B..E
#: hoja Carne, bloque A: (nombre, precio €/kg, útil, minutos de limpiar y picar por kg, % de grasa)
CARNES = [
    ('Blend comprado ya picado', 8.90, 1.00, 0.0, 0.20),
    ('Pieza para picar en el local (aguja o falda)', 7.50, 0.90, 6.0, 0.20),
    ('Bolas compradas porcionadas', 10.40, 1.00, 0.0, 0.20),
    ('Otra carne (escribe la tuya)', 9.50, 1.00, 0.0, 0.20),
]
P_ROW = OrderedDict([('pr', 6), ('ut', 7), ('tl', 8), ('gr', 9)])
P_NOTAS = OrderedDict([
    ('pr', 'Carne picada de vacuno: la «burger meat» minorista sale a 8,03 €/kg en oferta y el precio profesional no está publicado. El kit '
           'pone 8,90 €/kg el blend de aguja, falda y pecho con un 20 % de grasa, 7,50 €/kg la pieza para picar (sin fuente: mídelo) y '
           '10,40 €/kg las bolas porcionadas (sin fuente: pídelo a tu proveedor). Escribe el precio por kilo también de las bolas.' + EST),
    ('ut', 'Parte de la carne comprada que queda tras limpiar y picar. 100 % si la compras ya picada; 90 % en la pieza (sin fuente: mídelo): '
           'pesa la pieza antes y la carne limpia después.' + EST),
    ('tl', 'Minutos de limpiar y picar un kilo comprado. Sin fuente: mídelo. Cronometra 3 kilos y usa la media. Solo la pieza para picar '
           'lleva mano de obra.' + EST),
    ('gr', 'Informativo: no cambia ningún cálculo. El blend del plan lleva alrededor de un 20 % de grasa.' + EST),
])
OPC = [  # desplegable informativo: (opción, recordatorio)
    ('Compro la carne ya picada',
     'RD 1021/2022, art. 4.1: la carne picada se conserva a 2 °C o menos y la bola salada antes de la plancha es un preparado de carne: '
     '4 °C o menos. Pide al proveedor la ficha técnica con el blend, el % de grasa y la fecha de sacrificio o de envasado.'),
    ('Pico la carne en el local',
     'RD 1021/2022, arts. 7.2 y 7.3: la carne que picas en el local es del día y nunca de un día para otro; músculo sin recortes, con '
     'como máximo 6 días desde el sacrificio o 15 si va envasada al vacío. La carne madurada semanas queda fuera de ese plazo (Inf26008).'),
    ('Compro bolas porcionadas',
     'La bola salada antes de la plancha es un preparado de carne: 4 °C o menos (RD 1021/2022). Guarda la ficha técnica y la etiqueta con '
     'la fecha de cada lote.'),
]
#: bloque B (entradas comunes de la hoja Carne): (clave, rótulo, formato, valor, constante, nota)
ENT_C = [
    ('m', 'Merma de cierre (% de las bolas, el pan y las salsas que se tiran)', PCT, 0.03, False,
     'Las bolas que sobran al cierre y no se pueden guardar.' + EST),
    ('gb', 'Gramos por bola (g)', DEC, 100.0, False, 'Bola de 80-100 g; 100 g es la de Smash Hiro. Pesa 10 bolas y usa la media.' + EST),
    ('mc', 'Merma de cocción en plancha (% del peso de la bola)', PCT, 0.25, False,
     'Sin fuente: mídelo. Pesa 10 bolas antes de la plancha y después, y usa la media. Solo cambia los gramos servidos.' + EST),
    ('hh', 'Coste por hora del cocinero (€/h)', EUR, 16.0, False,
     '22.000 € brutos al año (lo más bajo de un cocinero de cadena, que gana de 22.000 a 25.000 €) + 33 % de cargas ÷ 1.840 h ≈ 16 €/h.' + EST),
    ('bc', 'Bolas por comensal', '#,##0.00', 1.6, False,
     'Criterio del plan, mídelo: con el 90 % de tus comensales en sencilla o doble y los menús, salen unas 1,6 bolas por comensal.' + EST),
    ('cd', 'Comensales equivalentes al día', ENT, 95, False,
     'El del caso del plan de negocio: 95 comensales equivalentes. Cámbialo por el tuyo.' + EST),
    ('cong', '¿Compras la carne picada o la picas en el local? (desplegable)', 'General', OPC[0][0], False,
     'Informativo: elige tu caso y mira el recordatorio de abajo. No cambia ningún cálculo.'),
    ('KG', 'Gramos por kilo (constante)', ENT, 1000, True, 'Constante: un kilo son 1.000 gramos. No se edita.'),
    ('MH', 'Minutos por hora (constante)', ENT, 60, True, 'Constante: una hora son 60 minutos. No se edita.'),
]
PE_M0 = 12
M_C = OrderedDict((k[0], PE_M0 + i) for i, k in enumerate(ENT_C))                  # 12..20
PE_RC = PE_M0 + len(ENT_C) + 1                                                     # 22: cabecera del bloque C
SAL_C = [
    ('cku', 'COSTE POR KG ÚTIL (€)', EUR, '{pr}/{ut}/(1-{m})', 'Precio ÷ útil ÷ (1 − merma de cierre).'),
    ('cku_mo', 'Coste por kg útil con la mano de obra de limpiar y picar (€)', EUR, '({pr}+{tl}/{MH}*{hh})/{ut}/(1-{m})',
     '(Precio + minutos ÷ 60 × coste por hora) ÷ útil ÷ (1 − merma de cierre). Alimenta el escandallo.'),
    ('c_bola', 'Coste de la carne de una bola (€)', EUR3, '{cku_mo}*{gb}/{KG}', 'Coste por kg útil con mano de obra × gramos ÷ 1.000.'),
    ('c_dos', 'Coste de la carne de una hamburguesa doble (2 bolas, €)', EUR3, '{c_bola}+{c_bola}', 'Dos bolas.'),
    ('c_tres', 'Coste de la carne de una hamburguesa triple (3 bolas, €)', EUR3, '{c_bola}+{c_bola}+{c_bola}', 'Tres bolas.'),
]
S_C = OrderedDict((k[0], PE_RC + 1 + i) for i, k in enumerate(SAL_C))             # 23..27
PE_RD = PE_RC + len(SAL_C) + 2                                                     # 29: cabecera del bloque D
SAL_D = [
    ('bkg', 'BOLAS POR KILO EN CRUDO (bolas)', DEC, '{KG}/{gb}', '1.000 ÷ gramos por bola. Es la cifra de la ficha técnica del proveedor.'),
    ('bkn', 'Bolas por kilo contando la merma de cierre (bolas)', DEC, '{KG}*(1-{m})/{gb}', '1.000 × (1 − merma de cierre) ÷ gramos por bola.'),
    ('gser', 'Gramos de carne que sirves tras la plancha (g)', DEC, '{gb}*(1-{mc})', 'Gramos por bola × (1 − merma de cocción).'),
    ('bdia', 'Bolas que vendes al día (bolas)', ENT, '{cd}*{bc}', 'Comensales equivalentes al día × bolas por comensal.'),
    ('carne_dia', 'KG DE CARNE AL DÍA', '#,##0.0', '{bdia}*{gb}/{KG}',
     'Bolas al día × gramos por bola ÷ 1.000. Dimensiona el pedido diario y el picado del día si picas en el local.'),
    ('lab_kg', 'Mano de obra de picar por kg útil de la pieza (€)', EUR, '{cku_mo1}-{cku1}',
     'Coste por kg útil de la pieza con mano de obra − sin ella.'),
    ('X', 'PICAR EN EL LOCAL TE COMPENSA MIENTRAS LIMPIARLO CUESTE MENOS DE (€ por kg útil)', EUR, '{cku_mo0}-{cku1}',
     'Coste por kg útil del blend comprado − coste por kg útil de la pieza (sin mano de obra).'),
    ('sem', 'Veredicto con tu mano de obra', 'General', 'IF({lab_kg}<{X},"%s","%s")' % (SEM_PICAR, SEM_BLEND),
     'Picar compensa si limpiarlo y picarlo cuesta menos que esa diferencia.'),
    ('rem', 'Recordatorio de tu caso', 'General', 'INDEX({rem_r},MATCH({cong},{opc_r},0))', 'Texto de la lista de abajo.'),
]
S_D = OrderedDict((k[0], PE_RD + 1 + i) for i, k in enumerate(SAL_D))             # 30..38
PE_RO = PE_RD + len(SAL_D) + 2                                                     # 40: cabecera de la lista
PE_O0 = PE_RO + 1                                                                  # 41..43
PE_O1 = PE_O0 + len(OPC) - 1

#: hoja Patatas: (clave, rótulo, formato, valor, constante, nota)
ENT_A = [
    ('pc', 'Patata congelada (€/kg, sin IVA)', EUR, 1.50, False,
     'Minorista a 1,76 €/kg (2,5 kg), McCain a 3,85 €/kg y el lineal a 2,10 €/kg de media (Circana, marzo de 2024); el kit pone 1,50 €/kg.' + EST),
    ('gc', 'Gramos de patata congelada por ración (g)', DEC, 150.0, False, 'Pesa 10 raciones y usa la media.' + EST),
    ('ac', 'Aceite que se queda en la ración congelada (% del peso, 1 g ≈ 1 ml)', PCT, 0.05, False, 'Lo que absorbe la patata al freírse.' + EST),
    ('pf', 'Patata fresca (€/kg, sin IVA)', EUR, 0.60, False, 'Sin fuente: precio profesional no publicado. Pídelo a tu proveedor.' + EST),
    ('gf', 'Gramos de patata fresca ya cortada por ración (g)', DEC, 150.0, False, 'Pesa 10 raciones y usa la media.' + EST),
    ('uf', 'Aprovechamiento de la patata fresca tras pelar y cortar (%)', PCT, 0.75, False,
     'Sin fuente: mídelo. Pesa 10 kg sin pelar y lo que queda cortado.' + EST),
    ('af', 'Aceite que se queda en la ración fresca (% del peso, 1 g ≈ 1 ml)', PCT, 0.10, False, 'La fresca absorbe más aceite que la congelada.' + EST),
    ('tf', 'Minutos de pelar, cortar y remojar un kilo de patata fresca (min)', DEC, 6.0, False,
     'Sin fuente: mídelo. Cronometra 3 kilos y usa la media.' + EST),
    ('hp', 'Coste por hora del cocinero (€/h)', EUR, 16.0, False, 'Igual que en la hoja Carne.' + EST),
    ('lc', 'Litros de aceite por cuba (l)', DEC, 10.0, False, 'Una cuba de freidora de 8 a 10 litros.' + EST),
    ('nc', 'Cubas de la freidora', ENT, 2, False, 'Freidora de dos cubas.' + EST),
    ('pa', 'Aceite de girasol alto oleico (€/l, sin IVA)', EUR, 1.80, False, 'Minorista a 2,25 €/l; en garrafa, el kit pone 1,80 €/l.' + EST),
    ('rcm', 'Raciones entre dos cambios de aceite', ENT, 900, False,
     'Las que aguante el aceite con los polares por debajo del 25 %. Sin fuente: mídelo con tus propias mediciones.' + EST),
    ('rpc', 'Raciones de patata por comensal', '#,##0.00', 0.8, False, 'Criterio del plan, mídelo: cuatro de cada cinco comensales piden patatas.' + EST),
    ('dm', 'Días de apertura al mes', ENT, 25, False, 'El caso del plan abre 300 días al año.' + EST),
    ('KG', 'Gramos por kilo (constante)', ENT, 1000, True, 'Constante: un kilo son 1.000 gramos. No se edita.'),
    ('ML', 'Mililitros por litro (constante)', ENT, 1000, True, 'Constante: un litro son 1.000 ml. No se edita.'),
    ('MH', 'Minutos por hora (constante)', ENT, 60, True, 'Constante: una hora son 60 minutos. No se edita.'),
]
SAL_A = [
    ('cd_c', 'Comensales equivalentes al día (de la hoja Carne)', ENT, '{cd_carne}', 'Se escribe en la hoja Carne.'),
    ('aceite_cambio', 'Coste del aceite de un cambio completo (€)', EUR, '{lc}*{nc}*{pa}', 'Litros por cuba × cubas × €/l.'),
    ('ace_cong', 'ACEITE POR RACIÓN CONGELADA (€)', EUR3, '{aceite_cambio}/{rcm}+{gc}*{ac}/{ML}*{pa}',
     'Cambio de aceite ÷ raciones entre cambios + aceite que se queda en la ración.'),
    ('pat_cong', 'Patata de una ración congelada (€)', EUR3, '{gc}*{pc}/{KG}', 'Gramos × €/kg ÷ 1.000.'),
    ('c_cong', 'COSTE POR RACIÓN CONGELADA (€)', EUR, '{pat_cong}+{ace_cong}', 'Patata + aceite.'),
    ('kg_bruto', 'Kilos de patata sin pelar por ración fresca (kg)', '#,##0.000', '{gf}/{uf}/{KG}', 'Gramos cortados ÷ aprovechamiento ÷ 1.000.'),
    ('pat_fre', 'Patata de una ración fresca (€)', EUR3, '{kg_bruto}*{pf}', 'Kilos sin pelar × €/kg.'),
    ('ace_fre', 'Aceite por ración fresca (€)', EUR3, '{aceite_cambio}/{rcm}+{gf}*{af}/{ML}*{pa}',
     'Cambio de aceite ÷ raciones entre cambios + aceite que se queda en la ración.'),
    ('c_fre', 'COSTE POR RACIÓN FRESCA, SIN MANO DE OBRA (€)', EUR, '{pat_fre}+{ace_fre}', 'Patata + aceite.'),
    ('mo_fre', 'Mano de obra de pelar, cortar y remojar una ración fresca (€)', EUR, '{tf}/{MH}*{hp}*{kg_bruto}',
     'Minutos por kilo ÷ 60 × coste por hora × kilos sin pelar.'),
    ('c_fre_mo', 'Coste por ración fresca con la mano de obra (€)', EUR, '{c_fre}+{mo_fre}', 'Alimenta el escandallo si eliges la fresca.'),
    ('mo_kg', 'Mano de obra por kilo de patata fresca sin pelar (€)', EUR, '{tf}/{MH}*{hp}', 'Minutos por kilo ÷ 60 × coste por hora.'),
    ('X', 'LA FRESCA TE COMPENSA MIENTRAS PELAR Y CORTAR CUESTE MENOS DE (€ por kg sin pelar)', EUR, '({c_cong}-{c_fre})/{kg_bruto}',
     '(Coste de la ración congelada − coste de la ración fresca sin mano de obra) ÷ kilos sin pelar por ración.'),
    ('sem', 'Veredicto con tu mano de obra', 'General', 'IF({mo_kg}<{X},"%s","%s")' % (SEM_FRE, SEM_CONG),
     'La fresca compensa si pelarla y cortarla cuesta menos que esa diferencia.'),
    ('rac_dia', 'RACIONES DE PATATA AL DÍA', ENT, '{cd_c}*{rpc}', 'Comensales equivalentes al día × raciones por comensal.'),
    ('kg_cong_dia', 'Kilos de patata congelada al día (kg)', '#,##0.0', '{rac_dia}*{gc}/{KG}', 'Raciones al día × gramos ÷ 1.000.'),
    ('kg_fre_dia', 'Kilos de patata fresca sin pelar al día (kg)', '#,##0.0', '{rac_dia}*{kg_bruto}', 'Raciones al día × kilos sin pelar por ración.'),
    ('l_mes', 'LITROS DE ACEITE AL MES (patata congelada)', ENT, '{rac_dia}*{dm}*({lc}*{nc}/{rcm}+{gc}*{ac}/{ML})',
     'Raciones al mes × (litros de un cambio ÷ raciones entre cambios + aceite que se queda en la ración).'),
]
M_A, M_RA, S_A = None, None, None


def _filas(ent, sal, h0=5):
    m = OrderedDict((k[0], h0 + 1 + i) for i, k in enumerate(ent))
    ra = h0 + 1 + len(ent) + 1
    s = OrderedDict((k[0], ra + 1 + i) for i, k in enumerate(sal))
    return m, ra, s


M_A, M_RA, S_A = _filas(ENT_A, SAL_A)

#: hoja Escandallo: datos generales
ESC_IN = [
    ('iva', 'IVA de la comida (%)', PCT, 0.10, 'Comida preparada, en sala, para llevar y a domicilio: 10 % (Ley 37/1992, art. 91).'),
    ('obj', 'Food cost objetivo (% sobre el precio sin IVA)', PCT, 0.28, 'El del plan de negocio: 28 %, el lado prudente.' + EST),
    ('tol', 'Margen de vigilancia sobre el objetivo (puntos)', PCT, 0.03, 'Hasta este margen por encima del objetivo, el semáforo dice «Vigila el coste».'),
    ('sel', 'Carne del escandallo (1 a 4: la columna de la hoja Carne)', '0', 1, '1 = blend comprado ya picado, 2 = pieza para picar, 3 = bolas porcionadas, 4 = libre.'),
    ('tp', 'Patata del escandallo (1 congelada, 2 fresca con mano de obra)', '0', 1, '1 = ración congelada, 2 = ración fresca con la mano de obra de pelar y cortar.'),
    ('pan', 'Pan de hamburguesa (€ por pan, sin IVA)', EUR, 0.50, 'Brioche minorista a 0,94 € la pieza; el kit pone 0,50 € en compra profesional.' + EST),
    ('qk', 'Queso cheddar en lonchas (€/kg, sin IVA)', EUR, 8.00, 'Minorista de 8,75 a 23 €/kg; el kit pone 8,00 €/kg.' + EST),
    ('gq', 'Gramos de queso por loncha (g)', DEC, 20.0, 'Una loncha de 20 g por bola. Pesa 10 lonchas y usa la media.' + EST),
    ('cl', 'Caja para llevar por pedido (€, sin IVA)', EUR, 0.45, 'Caja, bolsa y servilletas de un pedido para llevar.' + EST),
    ('cp', 'Caja de plataforma por pedido (€, sin IVA)', EUR, 0.60, 'Envase de reparto más sellado.' + EST),
    ('rec', 'Recargo de precio en la plataforma (%)', PCT, 0.15, 'Lo que subes el precio en la plataforma para compensar.' + EST),
    ('com', 'Comisión de la plataforma de reparto (%)', PCT, 0.30, 'Sobre el precio sin IVA del pedido. Revisa tu contrato.' + EST),
]
E = OrderedDict((k[0], 5 + i) for i, k in enumerate(ESC_IN))
E_BO, E_PA = 5 + len(ESC_IN), 5 + len(ESC_IN) + 1                                # enlaces a las hojas Carne y Patatas
E_CAB = E_PA + 2
#: producto: (nombre, bolas, lonchas de queso, panes, otros ingredientes €, raciones de patata, PVP con IVA de sala, ventas a la semana)
PRODS = [
    ('Hamburguesa simple (1 bola)', 1, 1, 1, 0.22, 0, 8.90, 140),
    ('Hamburguesa doble (2 bolas)', 2, 2, 1, 0.24, 0, 11.90, 240),
    ('Hamburguesa triple (3 bolas)', 3, 3, 1, 0.26, 0, 14.50, 40),
    ('Hamburguesa doble con bacon', 2, 2, 1, 0.51, 0, 12.90, 90),
    ('Hamburguesa estilo Oklahoma (2 bolas)', 2, 2, 1, 0.40, 0, 12.50, 50),
    ('Hamburguesa de la casa (2 bolas)', 2, 2, 1, 0.65, 0, 13.50, 70),
    ('Hamburguesa de pollo crujiente', 0, 1, 1, 1.75, 0, 10.90, 60),
    ('Hamburguesa vegetal', 0, 1, 1, 1.55, 0, 10.50, 20),
    ('Ración de patatas', 0, 0, 0, 0.0, 1, 3.90, 170),
    ('Ración de patatas con queso y bacon', 0, 0, 0, 0.64, 1, 6.50, 90),
    ('Batido', 0, 0, 0, 1.15, 0, 4.90, 60),
    ('Postre', 0, 0, 0, 1.00, 0, 4.50, 55),
    ('Menú: doble + patatas + refresco', 2, 2, 1, 0.69, 1, 15.90, 230),
    ('Pedido para dos: 2 dobles + 2 patatas', 4, 4, 2, 0.48, 2, 29.90, 35),
]
NPR = len(PRODS)
J_SIM, J_DOB, J_MENU = 0, 1, 12
E_P0 = E_CAB + 1
E_P1 = E_P0 + NPR - 1
E_T0 = E_P1 + 2
TOT = [
    ('ing', 'Ventas de la semana sin IVA (€)', EUR, 'SUMPRODUCT(%s,%s)'),
    ('cos', 'Coste de la comida vendida en la semana (€)', EUR, 'SUMPRODUCT(%s,%s)'),
    ('fcp', 'FOOD COST PONDERADO (%)', PCT, None),
    ('semp', 'Semáforo del food cost ponderado', 'General', None),
]
T = OrderedDict((k[0], E_T0 + 1 + i) for i, k in enumerate(TOT))
E_FIN = E_T0 + len(TOT) + 2
# columnas de la tabla
CN, CBO, CLQ, CPN, COT, CRA, CPVP, CUD = 'A', 'B', 'C', 'D', 'E', 'F', 'G', 'H'
CCO, CPS, CFC, CSE, CSG, CMS, CML, CMP, CRK = 'I', 'J', 'K', 'L', 'M', 'N', 'O', 'P', 'Q'
SALP = [('c_tot', CCO, EUR3), ('pvp_sin', CPS, EUR), ('fc', CFC, PCT), ('sem', CSE, 'General'), ('pvp_sug', CSG, EUR),
        ('sala', CMS, EUR), ('llevar', CML, EUR), ('plat', CMP, EUR), ('rank', CRK, ENT)]

EXPUESTOS = OrderedDict([
    ('COSTE_KG_UTIL_BLEND', (HC, 'B%d' % S_C['cku'])),
    ('COSTE_KG_UTIL_PIEZA', (HC, 'C%d' % S_C['cku'])),
    ('BOLAS_KG', (HC, 'B%d' % S_D['bkg'])),
    ('BOLAS_KG_NETAS', (HC, 'B%d' % S_D['bkn'])),
    ('CARNE_DIA_KG', (HC, 'B%d' % S_D['carne_dia'])),
    ('COSTE_SIMPLE', (HE, '%s%d' % (CCO, E_P0 + J_SIM))),
    ('COSTE_DOBLE', (HE, '%s%d' % (CCO, E_P0 + J_DOB))),
    ('FC_SIMPLE', (HE, '%s%d' % (CFC, E_P0 + J_SIM))),
    ('FC_DOBLE', (HE, '%s%d' % (CFC, E_P0 + J_DOB))),
    ('FC_PONDERADO', (HE, 'B%d' % T['fcp'])),
    ('COSTE_RACION_PATATA_CONG', (HP, 'B%d' % S_A['c_cong'])),
    ('COSTE_RACION_PATATA_FRESCA', (HP, 'B%d' % S_A['c_fre'])),
    ('ACEITE_RACION', (HP, 'B%d' % S_A['ace_cong'])),
    ('MARGEN_SALA_MENU', (HE, '%s%d' % (CMS, E_P0 + J_MENU))),
    ('MARGEN_PLATAFORMA_MENU', (HE, '%s%d' % (CMP, E_P0 + J_MENU))),
])


def valores_defecto():
    """Todas las entradas por defecto en un solo dict plano (claves = las de las celdas verdes)."""
    v = OrderedDict()
    for i, (_n, pr, ut, tl, gr) in enumerate(CARNES):
        v['pr%d' % i], v['ut%d' % i], v['tl%d' % i], v['gr%d' % i] = pr, ut, tl, gr
    for ent in (ENT_C, ENT_A):
        for k, _r, _f2, val, _c, _n in ent:
            v[k] = val
    for k, _r, _f2, val, _n in ESC_IN:
        v[k] = val
    for j, (_nom, bo, lq, pn, ot, ra, pvp, ud) in enumerate(PRODS):
        v['bo%d' % j], v['lq%d' % j], v['pn%d' % j], v['ot%d' % j] = bo, lq, pn, ot
        v['ra%d' % j], v['pvp%d' % j], v['ud%d' % j] = ra, pvp, ud
    return v


# ---------------------------------------------------------------------------------------------------------------------
# Cálculo en Python (espejo de las fórmulas) — fuente del autotest
# ---------------------------------------------------------------------------------------------------------------------
def calcular(v):
    r = OrderedDict()
    cols = []
    for i in range(NPE):
        c = OrderedDict()
        c['cku'] = v['pr%d' % i] / v['ut%d' % i] / (1 - v['m'])
        c['cku_mo'] = (v['pr%d' % i] + v['tl%d' % i] / v['MH'] * v['hh']) / v['ut%d' % i] / (1 - v['m'])
        c['c_bola'] = c['cku_mo'] * v['gb'] / v['KG']
        c['c_dos'] = c['c_bola'] + c['c_bola']
        c['c_tres'] = c['c_bola'] + c['c_bola'] + c['c_bola']
        cols.append(c)
    r['Pe'] = cols
    d = OrderedDict()
    d['bkg'] = v['KG'] / v['gb']
    d['bkn'] = v['KG'] * (1 - v['m']) / v['gb']
    d['gser'] = v['gb'] * (1 - v['mc'])
    d['bdia'] = v['cd'] * v['bc']
    d['carne_dia'] = d['bdia'] * v['gb'] / v['KG']
    d['lab_kg'] = cols[1]['cku_mo'] - cols[1]['cku']
    d['X'] = cols[0]['cku_mo'] - cols[1]['cku']
    d['sem'] = SEM_PICAR if d['lab_kg'] < d['X'] else SEM_BLEND
    d['rem'] = OPC[[o for o, _t in OPC].index(v['cong'])][1]
    r['D'] = d
    a = OrderedDict()
    a['cd_c'] = v['cd']
    a['aceite_cambio'] = v['lc'] * v['nc'] * v['pa']
    a['ace_cong'] = a['aceite_cambio'] / v['rcm'] + v['gc'] * v['ac'] / v['ML'] * v['pa']
    a['pat_cong'] = v['gc'] * v['pc'] / v['KG']
    a['c_cong'] = a['pat_cong'] + a['ace_cong']
    a['kg_bruto'] = v['gf'] / v['uf'] / v['KG']
    a['pat_fre'] = a['kg_bruto'] * v['pf']
    a['ace_fre'] = a['aceite_cambio'] / v['rcm'] + v['gf'] * v['af'] / v['ML'] * v['pa']
    a['c_fre'] = a['pat_fre'] + a['ace_fre']
    a['mo_fre'] = v['tf'] / v['MH'] * v['hp'] * a['kg_bruto']
    a['c_fre_mo'] = a['c_fre'] + a['mo_fre']
    a['mo_kg'] = v['tf'] / v['MH'] * v['hp']
    a['X'] = (a['c_cong'] - a['c_fre']) / a['kg_bruto']
    a['sem'] = SEM_FRE if a['mo_kg'] < a['X'] else SEM_CONG
    a['rac_dia'] = a['cd_c'] * v['rpc']
    a['kg_cong_dia'] = a['rac_dia'] * v['gc'] / v['KG']
    a['kg_fre_dia'] = a['rac_dia'] * a['kg_bruto']
    a['l_mes'] = a['rac_dia'] * v['dm'] * (v['lc'] * v['nc'] / v['rcm'] + v['gc'] * v['ac'] / v['ML'])
    r['A'] = a
    sel, tp = int(v['sel']), int(v['tp'])
    r['bola'] = cols[sel - 1]['c_bola']
    r['pata'] = a['c_cong'] if tp == 1 else a['c_fre_mo']
    ps = []
    for j in range(NPR):
        q = OrderedDict()
        q['c_tot'] = (v['bo%d' % j] * r['bola'] + v['pn%d' % j] * v['pan'] + v['lq%d' % j] * v['gq'] * v['qk'] / v['KG']
                      + v['ot%d' % j] + v['ra%d' % j] * r['pata'])
        q['pvp_sin'] = v['pvp%d' % j] / (1 + v['iva'])
        q['fc'] = q['c_tot'] / q['pvp_sin']
        q['sem'] = FC_OK if q['fc'] <= v['obj'] else (FC_VIG if q['fc'] <= v['obj'] + v['tol'] else FC_ALTO)
        q['pvp_sug'] = q['c_tot'] / v['obj'] * (1 + v['iva'])
        q['sala'] = q['pvp_sin'] - q['c_tot']
        q['llevar'] = q['sala'] - v['cl']
        q['plat'] = q['pvp_sin'] * (1 + v['rec']) * (1 - v['com']) - q['c_tot'] - v['cp']
        ps.append(q)
    for q in ps:
        q['rank'] = 1 + sum(1 for o in ps if o['sala'] > q['sala'])
    r['Pr'] = ps
    t = OrderedDict()
    t['ing'] = sum(q['pvp_sin'] * v['ud%d' % j] for j, q in enumerate(ps))
    t['cos'] = sum(q['c_tot'] * v['ud%d' % j] for j, q in enumerate(ps))
    t['fcp'] = 0 if t['ing'] == 0 else t['cos'] / t['ing']
    t['semp'] = FC_OK if t['fcp'] <= v['obj'] else (FC_VIG if t['fcp'] <= v['obj'] + v['tol'] else FC_ALTO)
    r['T'] = t
    return r


# ---------------------------------------------------------------------------------------------------------------------
# Helpers (de la Taquería, la Pollería y el Sushi Bar)
# ---------------------------------------------------------------------------------------------------------------------
def _f(ws, coord, v, b=False, sz=11, i=False, fmt=None, wrap=False, fill=None, locked=True):
    c = ws[coord]
    c.value = v
    c.font = Font(name=FUENTE, size=sz, bold=b, italic=i)
    if fmt:
        c.number_format = fmt
    c.alignment = Alignment(wrap_text=wrap, vertical='top')
    if fill:
        c.fill = fill
    c.protection = Protection(locked=locked)
    return c


def _nota(ws, fila, texto, desde='E', hasta='F', chars=80.0):
    if not texto:
        return
    _f(ws, '%s%d' % (desde, fila), texto, wrap=True)
    if desde != hasta:
        ws.merge_cells('%s%d:%s%d' % (desde, fila, hasta, fila))
    ws.row_dimensions[fila].height = max(ws.row_dimensions[fila].height or 0, 16, 15 * math.ceil(len(texto) / chars))


def _imprimir(ws, horizontal=False):
    ws.page_setup.paperSize = 9                                                   # A4
    ws.page_setup.orientation = 'landscape' if horizontal else 'portrait'
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.oddFooter.center.text = 'aichef.pro/%s · página &P de &N' % PID
    ws.print_options.horizontalCentered = True
    ws.protection.sheet = True                                                    # sin contraseña
    ws.sheet_view.showGridLines = False


def _dv(ws, tipo, celdas, f1=None, f2=None, op=None, titulo='Valor no válido', msg=''):
    dv = DataValidation(type=tipo, formula1=f1, formula2=f2, operator=op, allow_blank=False, showErrorMessage=True,
                        errorTitle=titulo, error=msg)
    ws.add_data_validation(dv)
    for c in celdas:
        dv.add(c)


def a_excel(expr, ref):
    return re.sub(r'\{(\w+)\}', lambda m: ref[m.group(1)], expr)


def _bloque(ws, titulo_in, ent, m, ra, sal, s, neg, chars=105.0, extra=None):
    """Escribe el bloque de entradas (verdes/grises) y el de resultados de una hoja de dos bloques."""
    _f(ws, 'A5', titulo_in, b=True, fill=GRIS)
    _f(ws, 'B5', 'Valor', b=True, fill=GRIS)
    _f(ws, 'C5', 'Notas', b=True, fill=GRIS)
    for k, rot, fmt, val, const, nota in ent:
        f = m[k]
        _f(ws, 'A%d' % f, rot)
        if const:
            _f(ws, 'B%d' % f, val, fmt=fmt, fill=GRIS)
        else:
            _f(ws, 'B%d' % f, val, fmt=fmt, fill=VERDE, locked=False)
        _nota(ws, f, nota, desde='C', hasta='I', chars=chars)
    ref = dict((k, '$B$%d' % f) for k, f in m.items())
    ref.update((k, '$B$%d' % f) for k, f in s.items())
    ref.update(extra or {})
    for col, t in zip('ABC', ('RESULTADOS', 'Valor', 'Notas')):
        _f(ws, '%s%d' % (col, ra), t, b=True, fill=GRIS)
    for k, rot, fmt, expr, nota in sal:
        f = s[k]
        _f(ws, 'A%d' % f, rot, b=k in neg)
        _f(ws, 'B%d' % f, '=' + a_excel(expr, ref), fmt=fmt, b=k in neg, wrap=(fmt == 'General'))
        _nota(ws, f, nota, desde='C', hasta='I', chars=chars)


def _semaforo(ws, celda, ok, ko):
    m = re.match(r'([A-Z]+)(\d+)', celda)
    ws.conditional_formatting.add(celda, FormulaRule(formula=['$%s$%s="%s"' % (m.group(1), m.group(2), ok)], fill=OK))
    ws.conditional_formatting.add(celda, FormulaRule(formula=['$%s$%s="%s"' % (m.group(1), m.group(2), ko)], fill=ROJO))


# ---------------------------------------------------------------------------------------------------------------------
# Hojas
# ---------------------------------------------------------------------------------------------------------------------
def hoja_carne(wb):
    ws = wb.active
    ws.title = HC
    for col, w in (('A', 62), ('B', 19), ('C', 19), ('D', 19), ('E', 19), ('F', 72)):
        ws.column_dimensions[col].width = w
    _f(ws, 'A1', 'CALCULADORA DE LA CARNE — COSTE POR KILO ÚTIL, POR BOLA, BOLAS POR KILO Y KILOS AL DÍA', b=True, sz=14)
    _f(ws, 'A2', 'Escribe solo en las celdas VERDES: todo lo demás se calcula. Importes sin IVA. Todos los valores por defecto son una '
                 'estimación del kit, editable: cámbialos por los de tu proveedor y, sobre todo, por lo que midas en tus propias bolas.', wrap=True)
    ws.merge_cells('A2:F2')
    ws.row_dimensions[2].height = 32
    _f(ws, 'A3', 'AVISO FIJO. Carne picada (RD 1021/2022, arts. 4.1, 7.2 y 7.3): comprada picada o picada en el local, se conserva a 2 °C o '
                 'menos; la bola salada antes de la plancha es un preparado de carne, a 4 °C o menos; la picada en el local es del día y nunca de '
                 'un día para otro. La carne madurada semanas queda fuera del plazo de picado de 6 días desde el sacrificio (15 si es vacuno al vacío) (Inf26008). Los minutos de limpiar y picar (6) y la merma de cocción '
                 'no tienen fuente: mídelos.', wrap=True, b=True, fill=GRIS)
    ws.merge_cells('A3:F3')
    ws.row_dimensions[3].height = 64
    # A. tres carnes y una libre
    _f(ws, 'A5', 'A. TRES CARNES Y UNA LIBRE', b=True, fill=GRIS)
    for i, c in enumerate(CARNES):
        _f(ws, '%s5' % PC[i], c[0], b=True, fill=VERDE, wrap=True, locked=False)
    _f(ws, 'F5', 'Notas', b=True, fill=GRIS)
    ws.row_dimensions[5].height = 62
    for k, rot, fmt, idx in (('pr', 'Precio de compra (€/kg, sin IVA)', EUR, 1), ('ut', 'Parte útil tras limpiar y picar (% del peso comprado)', PCT, 2),
                             ('tl', 'Minutos de limpiar y picar un kilo comprado (min)', DEC, 3),
                             ('gr', 'Grasa de la carne (%, informativo)', PCT, 4)):
        f = P_ROW[k]
        _f(ws, 'A%d' % f, rot)
        for i, p in enumerate(CARNES):
            _f(ws, '%s%d' % (PC[i], f), p[idx], fmt=fmt, fill=VERDE, locked=False)
        _nota(ws, f, P_NOTAS[k], desde='F', hasta='F', chars=75.0)
    # B. datos comunes
    _f(ws, 'A11', 'B. DATOS COMUNES', b=True, fill=GRIS)
    _f(ws, 'B11', 'Valor', b=True, fill=GRIS)
    _f(ws, 'C11', 'Notas', b=True, fill=GRIS)
    for k, rot, fmt, val, const, nota in ENT_C:
        f = M_C[k]
        _f(ws, 'A%d' % f, rot)
        if const:
            _f(ws, 'B%d' % f, val, fmt=fmt, fill=GRIS)
        else:
            _f(ws, 'B%d' % f, val, fmt=fmt, fill=VERDE, locked=False, wrap=(k == 'cong'))
        _nota(ws, f, nota, desde='C', hasta='F', chars=125.0)
    ws.merge_cells('B%d:E%d' % (M_C['cong'], M_C['cong']))
    ws.row_dimensions[M_C['cong']].height = 32
    # C. coste por kg útil
    _f(ws, 'A%d' % PE_RC, 'C. COSTE POR KG ÚTIL Y POR BOLA', b=True, fill=GRIS)
    for i in range(NPE):
        _f(ws, '%s%d' % (PC[i], PE_RC), '=%s5' % PC[i], b=True, fill=GRIS, wrap=True)
    _f(ws, 'F%d' % PE_RC, 'Notas', b=True, fill=GRIS)
    ws.row_dimensions[PE_RC].height = 62
    for k, rot, fmt, expr, nota in SAL_C:
        f = S_C[k]
        _f(ws, 'A%d' % f, rot, b=k in ('cku', 'cku_mo'))
        for i in range(NPE):
            c = PC[i]
            ref = dict((kk, '$B$%d' % ff) for kk, ff in M_C.items())
            ref.update((kk, '%s$%d' % (c, ff)) for kk, ff in P_ROW.items())
            ref.update((kk, '%s$%d' % (c, ff)) for kk, ff in S_C.items())
            _f(ws, '%s%d' % (c, f), '=' + a_excel(expr, ref), fmt=fmt, b=k in ('cku', 'cku_mo'))
        _nota(ws, f, nota, desde='F', hasta='F', chars=75.0)
    # D. bolas, kilos al día y picar o comprar
    _f(ws, 'A%d' % PE_RD, 'D. BOLAS POR KILO, KILOS AL DÍA Y ¿PICAR EN EL LOCAL?', b=True, fill=GRIS)
    _f(ws, 'B%d' % PE_RD, 'Valor', b=True, fill=GRIS)
    _f(ws, 'C%d' % PE_RD, 'Notas', b=True, fill=GRIS)
    ref = dict((k, '$B$%d' % f) for k, f in M_C.items())
    ref.update((k, '$B$%d' % f) for k, f in S_D.items())
    ref.update({'cku0': '$B$%d' % S_C['cku'], 'cku1': '$C$%d' % S_C['cku'], 'cku_mo0': '$B$%d' % S_C['cku_mo'],
                'cku_mo1': '$C$%d' % S_C['cku_mo'],
                'rem_r': '$B$%d:$B$%d' % (PE_O0, PE_O1), 'opc_r': '$A$%d:$A$%d' % (PE_O0, PE_O1)})
    neg = ('bkg', 'carne_dia', 'X', 'sem')
    for k, rot, fmt, expr, nota in SAL_D:
        f = S_D[k]
        _f(ws, 'A%d' % f, rot, b=k in neg, wrap=True)
        _f(ws, 'B%d' % f, '=' + a_excel(expr, ref), fmt=fmt, b=k in neg, wrap=(fmt == 'General'))
        _nota(ws, f, nota, desde='C', hasta='F', chars=125.0)
    ws.merge_cells('B%d:E%d' % (S_D['rem'], S_D['rem']))
    ws.row_dimensions[S_D['rem']].height = 48
    ws.row_dimensions[S_D['X']].height = 32
    ws.row_dimensions[S_D['sem']].height = 32
    _semaforo(ws, 'B%d' % S_D['sem'], SEM_PICAR, SEM_BLEND)
    # lista del desplegable
    _f(ws, 'A%d' % PE_RO, 'LISTA DEL DESPLEGABLE DE CARNE PICADA', b=True, fill=GRIS)
    _f(ws, 'B%d' % PE_RO, 'Recordatorio', b=True, fill=GRIS)
    for i, (op, rem) in enumerate(OPC):
        f = PE_O0 + i
        _f(ws, 'A%d' % f, op, wrap=True)
        _f(ws, 'B%d' % f, rem, wrap=True)
        ws.merge_cells('B%d:F%d' % (f, f))
        ws.row_dimensions[f].height = 32
    # validaciones
    _dv(ws, 'decimal', ['%s%d' % (c, P_ROW['pr']) for c in PC], f1='0', op='greaterThan', msg='Escribe un número mayor que 0.')
    _dv(ws, 'decimal', ['%s%d' % (c, P_ROW['ut']) for c in PC], f1='0.3', f2='1', op='between', titulo='Porcentaje no válido',
        msg='Escribe entre 0,3 y 1 (0,9 = 90 %).')
    _dv(ws, 'decimal', ['%s%d' % (c, P_ROW['tl']) for c in PC], f1='0', op='greaterThanOrEqual', msg='Escribe un número mayor o igual que 0.')
    _dv(ws, 'decimal', ['%s%d' % (c, P_ROW['gr']) for c in PC] + ['B%d' % M_C['m'], 'B%d' % M_C['mc']], f1='0', f2='0.5', op='between',
        titulo='Porcentaje no válido', msg='Escribe un porcentaje entre 0 y 0,5 (0,05 = 5 %).')
    _dv(ws, 'decimal', ['B%d' % M_C[k] for k in ('gb', 'bc')], f1='0', op='greaterThan', msg='Escribe un número mayor que 0.')
    _dv(ws, 'decimal', ['B%d' % M_C['hh']], f1='0', op='greaterThanOrEqual', msg='Escribe un número mayor o igual que 0.')
    _dv(ws, 'whole', ['B%d' % M_C['cd']], f1='1', f2='1000000', op='between', titulo='Número no válido', msg='Escribe un número entero de 1 o más.')
    _dv(ws, 'list', ['B%d' % M_C['cong']], f1='$A$%d:$A$%d' % (PE_O0, PE_O1), titulo='Opción no válida', msg='Elige una opción de la lista.')
    _imprimir(ws, horizontal=True)


def hoja_patatas(wb):
    ws = wb.create_sheet(HP)
    for col, w in (('A', 66), ('B', 15), ('C', 15), ('D', 16), ('E', 14), ('F', 14), ('G', 14), ('H', 12), ('I', 15)):
        ws.column_dimensions[col].width = w
    _f(ws, 'A1', 'CALCULADORA DE LAS PATATAS Y EL ACEITE — COSTE POR RACIÓN, KILOS AL DÍA Y LITROS AL MES', b=True, sz=14)
    _f(ws, 'A2', 'Escribe solo en las celdas VERDES: todo lo demás se calcula. Importes sin IVA. Todos los valores por defecto son una '
                 'estimación del kit, editable: pesa tus raciones, cronometra el pelado y cambia los datos por los tuyos.', wrap=True)
    ws.merge_cells('A2:I2')
    ws.row_dimensions[2].height = 32
    _f(ws, 'A3', 'AVISO FIJO. Acrilamida (Reglamento (UE) 2017/2158): fríe a menos de 175 °C, guarda la patata fresca a más de 6 °C, remójala o '
                 'escáldala antes de freírla y compara el color con la guía de colores de patatas. Aceite (Orden de 26 de enero de 1989): polares '
                 'por debajo del 25 %, nada ajeno al aceite añadido al baño y el usado, a un gestor autorizado; la norma no obliga a tener un medidor.',
       wrap=True, b=True, fill=GRIS)
    ws.merge_cells('A3:I3')
    ws.row_dimensions[3].height = 50
    _bloque(ws, 'A. LA PATATA Y EL ACEITE', ENT_A, M_A, M_RA, SAL_A, S_A, ('c_cong', 'c_fre', 'X', 'sem', 'rac_dia', 'l_mes', 'ace_cong'),
            extra={'cd_carne': "'%s'!$B$%d" % (HC, M_C['cd'])})
    ws.row_dimensions[S_A['X']].height = 32
    ws.row_dimensions[S_A['sem']].height = 32
    _semaforo(ws, 'B%d' % S_A['sem'], SEM_FRE, SEM_CONG)
    cel = lambda k: 'B%d' % M_A[k]
    _dv(ws, 'decimal', [cel(k) for k in ('pc', 'gc', 'pf', 'gf', 'hp', 'lc', 'pa', 'rcm', 'rpc')], f1='0', op='greaterThan',
        msg='Escribe un número mayor que 0.')
    _dv(ws, 'decimal', [cel('uf')], f1='0.3', f2='1', op='between', titulo='Porcentaje no válido', msg='Escribe entre 0,3 y 1 (0,75 = 75 %).')
    _dv(ws, 'decimal', [cel('tf')], f1='0', op='greaterThanOrEqual', msg='Escribe un número mayor o igual que 0.')
    _dv(ws, 'decimal', [cel('ac'), cel('af')], f1='0', f2='0.5', op='between', titulo='Porcentaje no válido',
        msg='Escribe un porcentaje entre 0 y 0,5 (0,05 = 5 %).')
    _dv(ws, 'whole', [cel('nc')], f1='1', f2='20', op='between', titulo='Número no válido', msg='Escribe un número entero de 1 a 20.')
    _dv(ws, 'whole', [cel('dm')], f1='1', f2='31', op='between', titulo='Número no válido', msg='Escribe un número entero de 1 a 31.')
    _imprimir(ws, horizontal=True)


def hoja_escandallo(wb):
    ws = wb.create_sheet(HE)
    for col, w in ((CN, 44), (CBO, 11), (CLQ, 11), (CPN, 10), (COT, 13), (CRA, 12), (CPVP, 13), (CUD, 14), (CCO, 12), (CPS, 12), (CFC, 11),
                   (CSE, 17), (CSG, 14), (CMS, 12), (CML, 12), (CMP, 13), (CRK, 10)):
        ws.column_dimensions[col].width = w
    _f(ws, 'A1', 'ESCANDALLO DE LA HAMBURGUESERÍA — 14 PRODUCTOS, POR PRODUCTO Y POR CANAL', b=True, sz=14)
    _f(ws, 'A2', 'Escribe solo en las celdas VERDES (nombres, bolas, queso, pan, otros ingredientes, patatas, precios de carta y ventas). '
                 'Importes sin IVA salvo el precio de carta. Todos los valores por defecto son una estimación del kit, editable: pon los tuyos. '
                 'El coste de la bola viene de la hoja «%s», el de la ración de patata de la hoja «%s».' % (HC, HP), wrap=True)
    ws.merge_cells('A2:%s2' % CRK)
    ws.row_dimensions[2].height = 32
    for col, t in zip('ABD', ('DATOS GENERALES', 'Valor', 'Notas')):
        _f(ws, '%s4' % col, t, b=True, fill=GRIS)
    for k, rot, fmt, val, nota in ESC_IN:
        f = E[k]
        _f(ws, 'A%d' % f, rot)
        _f(ws, 'B%d' % f, val, fmt=fmt, fill=VERDE, locked=False)
        _nota(ws, f, nota, desde='D', hasta=CRK, chars=150.0)
    _f(ws, 'A%d' % E_BO, 'Coste de la carne de una bola (€, de la hoja Carne)')
    _f(ws, 'B%d' % E_BO, "=CHOOSE($B$%d,%s)" % (E['sel'], ','.join("'%s'!$%s$%d" % (HC, c, S_C['c_bola']) for c in PC)), fmt=EUR3)
    _f(ws, 'A%d' % E_PA, 'Coste de una ración de patata (€, de la hoja Patatas)')
    _f(ws, 'B%d' % E_PA, "=CHOOSE($B$%d,'%s'!$B$%d,'%s'!$B$%d)" % (E['tp'], HP, S_A['c_cong'], HP, S_A['c_fre_mo']), fmt=EUR3)
    cab = [(CN, 'PRODUCTO'), (CBO, 'Bolas de carne (nº)'), (CLQ, 'Lonchas de queso (nº)'), (CPN, 'Panes (nº)'), (COT, 'Otros ingredientes (€)'),
           (CRA, 'Raciones de patata (nº)'), (CPVP, 'PVP con IVA de sala (€)'), (CUD, 'Vendidos a la semana (nº)'), (CCO, 'COSTE (€)'),
           (CPS, 'PVP sin IVA (€)'), (CFC, 'FOOD COST (%)'), (CSE, 'Semáforo'), (CSG, 'PVP con IVA sugerido (€)'),
           (CMS, 'Margen en sala (€)'), (CML, 'Margen para llevar (€)'), (CMP, 'Margen en plataforma (€)'), (CRK, 'Puesto por margen')]
    for col, t in cab:
        _f(ws, '%s%d' % (col, E_CAB), t, b=True, fill=GRIS, wrap=True)
    ws.row_dimensions[E_CAB].height = 48
    kg = "'%s'!$B$%d" % (HC, M_C['KG'])
    for j, (nom, bo, lq, pn, ot, ra, pvp, ud) in enumerate(PRODS):
        f = E_P0 + j
        _f(ws, 'A%d' % f, nom, fill=VERDE, locked=False, b=True)
        for col, val, fmt in ((CBO, bo, '#,##0.0;-#,##0.0;"-"'), (CLQ, lq, '#,##0.0;-#,##0.0;"-"'), (CPN, pn, '#,##0.0;-#,##0.0;"-"'),
                              (COT, ot, EUR), (CRA, ra, '#,##0.0;-#,##0.0;"-"'), (CUD, ud, ENT)):
            _f(ws, '%s%d' % (col, f), float(val), fmt=fmt, fill=VERDE, locked=False)
        _f(ws, '%s%d' % (CPVP, f), float(pvp), fmt=EUR, fill=VERDE, locked=False, b=True)
        ex = {
            CCO: '%s%d*$B$%d+%s%d*$B$%d+%s%d*$B$%d*$B$%d/%s+%s%d+%s%d*$B$%d' % (
                CBO, f, E_BO, CPN, f, E['pan'], CLQ, f, E['gq'], E['qk'], kg, COT, f, CRA, f, E_PA),
            CPS: '%s%d/(1+$B$%d)' % (CPVP, f, E['iva']),
            CFC: '%s%d/%s%d' % (CCO, f, CPS, f),
            CSE: 'IF(%s%d<=$B$%d,"%s",IF(%s%d<=$B$%d+$B$%d,"%s","%s"))' % (
                CFC, f, E['obj'], FC_OK, CFC, f, E['obj'], E['tol'], FC_VIG, FC_ALTO),
            CSG: '%s%d/$B$%d*(1+$B$%d)' % (CCO, f, E['obj'], E['iva']),
            CMS: '%s%d-%s%d' % (CPS, f, CCO, f),
            CML: '%s%d-$B$%d' % (CMS, f, E['cl']),
            CMP: '%s%d*(1+$B$%d)*(1-$B$%d)-%s%d-$B$%d' % (CPS, f, E['rec'], E['com'], CCO, f, E['cp']),
            CRK: '1+COUNTIF($%s$%d:$%s$%d,">"&%s%d)' % (CMS, E_P0, CMS, E_P1, CMS, f),
        }
        for _k, col, fmt in SALP:
            neg = col in (CCO, CFC)
            _f(ws, '%s%d' % (col, f), '=' + ex[col], fmt=fmt, b=neg or col == CRK, wrap=(col == CSE))
        ws.row_dimensions[f].height = 32
    rng = '%s%d:%s%d' % (CSE, E_P0, CSE, E_P1)
    top = '%s%d' % (CSE, E_P0)
    for txt, fill in ((FC_OK, OK), (FC_VIG, AMBAR), (FC_ALTO, ROJO)):
        ws.conditional_formatting.add(rng, FormulaRule(formula=['%s="%s"' % (top, txt)], fill=fill))
    # ponderado
    _f(ws, 'A%d' % E_T0, 'FOOD COST PONDERADO DE LA SEMANA (suma de costes × ventas ÷ suma de PVP sin IVA × ventas)', b=True, fill=GRIS)
    _f(ws, 'B%d' % E_T0, 'Valor', b=True, fill=GRIS)
    rc = '$%s$%d:$%s$%d' % (CCO, E_P0, CCO, E_P1)
    rp = '$%s$%d:$%s$%d' % (CPS, E_P0, CPS, E_P1)
    ru = '$%s$%d:$%s$%d' % (CUD, E_P0, CUD, E_P1)
    exs = {'ing': 'SUMPRODUCT(%s,%s)' % (rp, ru), 'cos': 'SUMPRODUCT(%s,%s)' % (rc, ru),
           'fcp': 'IF($B$%d=0,0,$B$%d/$B$%d)' % (T['ing'], T['cos'], T['ing']),
           'semp': 'IF($B$%d<=$B$%d,"%s",IF($B$%d<=$B$%d+$B$%d,"%s","%s"))' % (
               T['fcp'], E['obj'], FC_OK, T['fcp'], E['obj'], E['tol'], FC_VIG, FC_ALTO)}
    for k, rot, fmt, _x in TOT:
        f = T[k]
        _f(ws, 'A%d' % f, rot, b=k == 'fcp')
        _f(ws, 'B%d' % f, '=' + exs[k], fmt=fmt, b=k == 'fcp', wrap=(k == 'semp'))
    ws.row_dimensions[T['semp']].height = 32
    for txt, fill in ((FC_OK, OK), (FC_VIG, AMBAR), (FC_ALTO, ROJO)):
        ws.conditional_formatting.add('B%d' % T['semp'], FormulaRule(formula=['$B$%d="%s"' % (T['semp'], txt)], fill=fill))
    _f(ws, 'A%d' % E_FIN, 'El food cost de la tabla es el del producto servido en sala. En el pedido para llevar se descuenta la caja; en la '
                          'plataforma, además, se sube el precio con el recargo y se descuenta la comisión sobre ese precio. El margen se ordena por '
                          'el de sala. AVISO FIJO: las patatas viajan mal; en la plataforma, envase ventilado o déjalas fuera del pedido de reparto. '
                          'El menú se calcula con un solo IVA, el de la comida; si vendes mucha bebida, separa las dos bases. Este coste es TEÓRICO: '
                          'el plan de negocio trabaja con un 28 %, el lado prudente; la diferencia entre la receta y la caja está explicada en la '
                          'hoja «Instrucciones».', i=True, wrap=True)
    ws.merge_cells('A%d:%s%d' % (E_FIN, CRK, E_FIN))
    ws.row_dimensions[E_FIN].height = 66
    ws.freeze_panes = 'B%d' % (E_CAB + 1)
    # validaciones
    _dv(ws, 'decimal', ['B%d' % E[k] for k in ('iva', 'tol', 'rec', 'com')], f1='0', f2='0.9', op='between',
        titulo='Porcentaje no válido', msg='Escribe un porcentaje entre 0 y 0,9 (0,1 = 10 %).')
    _dv(ws, 'decimal', ['B%d' % E['obj']], f1='0.01', f2='0.9', op='between', titulo='Objetivo no válido',
        msg='El objetivo no puede ser 0 (se divide entre él): escribe entre 0,01 y 0,9.')
    _dv(ws, 'decimal', ['B%d' % E[k] for k in ('pan', 'qk', 'gq', 'cl', 'cp')]
        + ['%s%d' % (c, f) for c in (CBO, CLQ, CPN, COT, CRA, CPVP, CUD) for f in range(E_P0, E_P1 + 1)],
        f1='0', op='greaterThanOrEqual', titulo='Importe no válido', msg='Escribe un número mayor o igual que 0.')
    _dv(ws, 'whole', ['B%d' % E['sel']], f1='1', f2=str(NPE), op='between', titulo='Carne no válida',
        msg='Escribe 1, 2, 3 o 4, el número de la columna de la hoja Carne.')
    _dv(ws, 'whole', ['B%d' % E['tp']], f1='1', f2='2', op='between', titulo='Patata no válida', msg='Escribe 1 (congelada) o 2 (fresca).')
    _imprimir(ws, horizontal=True)


def _n(x, d=2):
    """Número en formato español (coma decimal, punto de millares)."""
    s = ('{:,.%df}' % d).format(x)
    return s.replace(',', '§').replace('.', ',').replace('§', '.')


def instrucciones():
    r = calcular(valores_defecto())
    pe, d, a, pr, t = r['Pe'], r['D'], r['A'], r['Pr'], r['T']
    menu = pr[J_MENU]
    return [
        ('INSTRUCCIONES DE USO — Calculadora de la carne, las patatas y el escandallo', True),
        ('Tres herramientas en un libro: la hoja «Carne» te dice cuánto cuesta el kilo útil del blend comprado ya picado, de la pieza para picar en '
         'el local y de las bolas porcionadas, cuánto vale cada bola, cuántas bolas salen de un kilo, cuántos kilos de carne necesitas al día y si '
         'te compensa picar en el local; la hoja «Patatas» calcula el coste de la ración congelada y de la fresca, el aceite por ración, los kilos de '
         'patata al día y los litros de aceite al mes; la hoja «Escandallo» calcula el coste, el food cost, el margen y el precio de carta de 14 '
         'productos, con el margen en sala, para llevar y en plataforma, y el food cost ponderado de tu semana. Solo se escribe en las celdas '
         'VERDES; lo demás se calcula.', False),
        ('Todos los valores por defecto son estimación del kit, editable: son ejemplos coherentes con el caso del plan de negocio, no precios de '
         'nadie. Cámbialos por los de tu proveedor y por lo que midas en tu plancha y tu freidora antes de decidir precios.', False),
        ('HOJA «CARNE»', True),
        ('1. Pesa 10 bolas y usa la media en los gramos por bola. Pesa también 10 bolas antes y después de la plancha para medir la merma de cocción '
         '(sin fuente: mídelo). Si picas en el local, pesa la pieza antes de limpiarla y la carne limpia después, y cronometra cuánto tardas en '
         'limpiar y picar un kilo.', False),
        ('2. El coste por kilo útil es el precio ÷ útil ÷ (1 − merma de cierre). Con los valores del kit sale a %s €/kg el blend comprado ya picado y '
         'a %s €/kg la pieza para picar, sin contar la mano de obra. El coste de una bola es ese coste por kilo útil con la mano de obra × los '
         'gramos ÷ 1.000: %s € con el blend.' % (_n(pe[0]['cku']), _n(pe[1]['cku']), _n(pe[0]['c_bola'], 3)), False),
        ('3. Bolas por kilo. En crudo son 1.000 ÷ gramos por bola: %s bolas de 100 g por kilo, la cifra de la ficha técnica del proveedor. Contando '
         'la merma de cierre del 3 %% salen %s. Con 1,6 bolas por comensal y %s comensales equivalentes al día, necesitas unos %s kg de carne al día: '
         'es el pedido diario y, si picas en el local, el picado del día.' % (_n(d['bkg'], 0), _n(d['bkn'], 1), _n(95, 0), _n(d['carne_dia'], 1)), False),
        ('4. ¿Picar o comprar? Con los valores del kit, picar en el local solo te compensa mientras limpiar y picar un kilo cueste menos de %s € por '
         'kilo útil; con 6 minutos por kilo y 16 € la hora cuesta %s €: el veredicto es «%s». Mide tus minutos antes de comprar la picadora.'
         % (_n(d['X']), _n(d['lab_kg']), d['sem']), False),
        ('5. Carne picada (RD 1021/2022, arts. 4.1, 7.2 y 7.3). Comprada picada o picada en el local, se conserva a 2 °C o menos; la bola salada '
         'antes de la plancha es un preparado de carne, a 4 °C o menos; la picada en el local es del día y nunca de un día para otro, con músculo '
         'sin recortes y como máximo 6 días desde el sacrificio, o 15 si va envasada al vacío. La carne madurada semanas queda fuera de ese plazo (Inf26008). '
         'El desplegable de la hoja es solo un recordatorio: no cambia ningún cálculo.', False),
        ('HOJA «PATATAS»', True),
        ('1. Pesa 10 raciones de patata y usa la media. Si usas fresca, pesa 10 kg sin pelar y lo que queda cortado para medir el aprovechamiento, y '
         'cronometra el pelado, el corte y el remojo de un kilo. Para el aceite, anota cuántas raciones aguanta cada cambio con los polares por '
         'debajo del 25 % (sin fuente: mídelo).', False),
        ('2. Con los valores del kit, una ración congelada cuesta %s € (de ellos, %s € de aceite) y una fresca %s € sin mano de obra y %s € con ella. '
         'La fresca solo te compensa mientras pelar y cortar un kilo cueste menos de %s €; con 6 minutos por kilo y 16 € la hora cuesta %s €: '
         'el veredicto es «%s».' % (_n(a['c_cong']), _n(a['ace_cong']), _n(a['c_fre']), _n(a['c_fre_mo']), _n(a['X']), _n(a['mo_kg']), a['sem']), False),
        ('3. Con %s raciones al día (comensales × 0,8) necesitas unos %s kg de patata congelada al día, o %s kg de fresca sin pelar, y unos %s litros '
         'de aceite al mes.' % (_n(a['rac_dia'], 0), _n(a['kg_cong_dia'], 1), _n(a['kg_fre_dia'], 1), _n(a['l_mes'], 0)), False),
        ('4. Acrilamida (Reglamento (UE) 2017/2158): fríe a menos de 175 °C, guarda la patata fresca a más de 6 °C, remójala o escáldala antes de '
         'freírla y compara el color con la guía de colores de patatas. Aceite (Orden de 26 de enero de 1989): polares por debajo del 25 %, nada '
         'ajeno al aceite añadido al baño y el usado, a un gestor autorizado; la norma no obliga a tener un medidor, pero conviene medir.', False),
        ('HOJA «ESCANDALLO»', True),
        ('1. Escribe el IVA, el food cost objetivo, el margen de vigilancia, qué carne y qué patata usas (los números de las hojas anteriores), el '
         'pan, el queso, las cajas de para llevar y de plataforma, el recargo y la comisión de la plataforma.', False),
        ('2. En la tabla, cada fila es un producto. Escribe las bolas, las lonchas de queso (20 g cada una, una por bola), los panes, lo que cuesta '
         'el resto de ingredientes (salsa, encurtido, cebolla, mantequilla del pan y papel), las raciones de patata, el precio de carta con IVA y '
         'los productos que vendes a la semana.', False),
        ('3. Para cada producto ves el coste, el precio sin IVA, el food cost con semáforo, el precio de carta con IVA que necesitarías para llegar a '
         'tu objetivo y el margen en sala, para llevar y en plataforma, con el puesto de cada producto por margen en sala. Abajo, el food cost '
         'ponderado de tu semana. Con los valores del kit, la hamburguesa doble (2 bolas) cuesta %s € y tiene un food cost del %s %%; el menú '
         'deja %s € en sala, %s € para llevar y %s € en plataforma: en plataforma el margen cae casi un tercio.'
         % (_n(pr[J_DOB]['c_tot']), _n(pr[J_DOB]['fc'] * 100, 1), _n(menu['sala']), _n(menu['llevar']), _n(menu['plat'])), False),
        ('DEL FOOD COST TEÓRICO AL 28 % DEL PLAN', True),
        ('El food cost de esta hoja es teórico: el que saldría si todo lo que compras se vendiera tal como está en el escandallo. Con los valores '
         'del kit, la semana ponderada sale en un %s %%. El plan de negocio trabaja con un 28 %%, el lado prudente: entre la receta y la caja se '
         'pierden las bolas, el pan y las salsas que se tiran al cierre, las comidas del personal, las invitaciones, los errores y las reclamaciones '
         'de plataforma y las promociones de plataforma (2 × 1, primera compra); en el plan suman unos 4 puntos. Cada euro por kilo que sube la '
         'carne pone unos %s € en la hamburguesa doble (cerca de 2 puntos de food cost).'
         % (_n(t['fcp'] * 100, 1), _n(2 * 100 / 1000.0 / (1 - 0.03))), False),
        ('Tu comparación real es esta: las compras de comida del mes ÷ las ventas de comida sin IVA del mes. Si después del primer trimestre sale '
         'por encima o por debajo del 28 %, cambia la celda de coste de mercancía sobre las ventas de comida del plan.', False),
        ('PRECIOS, IVA Y FUENTES', True),
        ('El IVA de la comida preparada para consumir en el local, para llevar o a domicilio es el 10 % (Ley 37/1992, art. 91). La cerveza y los '
         'refrescos con azúcares o edulcorantes añadidos que salen del local, también los zero, van al 21 %: esta calculadora usa un solo IVA, el de la comida. Los '
         'importes van sin IVA, salvo el precio de carta.', False),
        ('Cocción. La norma no fija una cifra para la hamburguesa: la AESAN recomienda 70 °C o más durante al menos 2 minutos, tu APPCC decide el '
         'punto crítico con una sonda y la smash fina sale hecha por su grosor.', False),
        ('Los precios y rendimientos salen del estudio de mercado del kit (octubre de 2026): carne picada y patata de minorista y mayorista, queso y '
         'pan de minorista y la carta de varias hamburgueserías de smash como referencia de precios. Ninguna cifra de esta calculadora es un dato '
         'publicado de tu negocio: sustitúyelas por las tuyas.', False),
    ]


def hoja_instrucciones(wb):
    ws = wb.create_sheet(HI)
    ws.column_dimensions['A'].width = 110
    f = 1
    for t, b in instrucciones():
        _f(ws, 'A%d' % f, t, b=b, sz=13 if f == 1 else 11, wrap=True)
        if not b:
            ws.row_dimensions[f].height = max(16, 15 * math.ceil(len(t) / 120.0))
        f += 2 if b and f > 1 else 1
    f += 1
    for t in ('Para editar una celda que no esté en verde: Revisar → Desproteger hoja (no tiene contraseña).',
              'Más plantillas y kits del catálogo en aichef.pro/productos-digitales', motor.BIO_LINE,
              motor.VERSION_FMT.format(pid=PID)):
        _f(ws, 'A%d' % f, t, wrap=True)
        f += 1
    _imprimir(ws)


def generar(destino):
    wb = openpyxl.Workbook()
    hoja_carne(wb)
    hoja_patatas(wb)
    hoja_escandallo(wb)
    hoja_instrucciones(wb)
    pr = wb.properties
    pr.title = '%s · v%s' % (TITULO, motor.VERSION)
    pr.subject = 'Plan de Negocio: Hamburguesería Smash'
    pr.description = 'aichef.pro/%s' % PID
    pr.keywords = 'plan de negocio hamburguesería, calculadora de carne y patatas, escandallo de hamburguesas, smash burger, AI Chef Pro'
    pr.creator = pr.lastModifiedBy = motor.CREATOR
    os.makedirs(os.path.dirname(destino), exist_ok=True)
    wb.save(destino)
    sys.path.insert(0, PD)
    import inject_cache as IC                                                      # noqa: E402
    nf, inj, fail = IC.inject(destino)
    print('%s: formulas=%d cache_inyectado=%d fallos_pycel=%d' % (os.path.basename(destino), nf, inj, fail))
    return fail


def valores_expuestos(ruta):
    """Lee de la CACHÉ del xlsx los valores expuestos: {nombre: valor}."""
    wv = openpyxl.load_workbook(ruta, data_only=True)
    return OrderedDict((n, wv[h][c].value) for n, (h, c) in EXPUESTOS.items())


# ---------------------------------------------------------------------------------------------------------------------
# Autotest
# ---------------------------------------------------------------------------------------------------------------------
#: ERROR del barrido (F1-camino §4): restos del Sushi Bar, la Taquería, la Pollería y la Pizzería + errores legales.
#: «netas» queda fuera de los textos de la calculadora (se dice «con la merma de cierre»).
PROHIBIDOS = [r'sushi', r'sashimi', r'nigiri', r'\bmakis?\b', r'uramaki', r'temaki', r'\bpoke\b', r'itamae', r'sushiman', r'\bshari\b', r'\bnori\b',
              r'\bnetas?\b', r'wasabi', r'\bsake\b', r'gyozas?', r'edamame', r'mochi', r'tempura', r'anisakis', r'hangiri', r'arroceras?', r'ph-metro',
              r'japon[eé]s', r'combinados?', r'salm[oó]n', r'at[uú]n', r'surimi', r'vitrina refrigerada de sushi', r'−20 °C', r'1420/2006',
              r'plan-negocio-sushi-bar',
              r'taquer', r'trompo', r'\btacos?\b', r'comal', r'al pastor', r'gringa', r'nixtamal', r'mezcal', r'tequila', r'aguas? frescas?', r'carnitas',
              r'suadero', r'cochinita', r'guacamole', r'tortillas? de', r'plan-negocio-taqueria',
              r'poller', r'asador de pollos', r'pollo a la brasa', r'pollo asado', r'broaster',
              r'pizz', r'napol[ei]t', r'mozzarella', r'\ble[ñn]a\b',
              r'validad[oa]', r'\btapas?\b',
              r'(?<!\d)75 ?°C', r'155 ?°F|(?<!\d)68 ?°C', r'temperatura legal', r'obligatoria']
AVISOS = [r'−18 °C', r'madurad', r'al punto|poco hecha', r'sin gluten', r'auténtic', r'gourmet', r'Día de la Hamburguesa', r'digital obligatorio',
          r'180 °C|190 °C', r'pan común', r'recongel', r'medidor']
CELDA = {}                                                                         # clave de entrada → (hoja, celda)
for _i in range(NPE):
    for _k in ('pr', 'ut', 'tl', 'gr'):
        CELDA['%s%d' % (_k, _i)] = (HC, '%s%d' % (PC[_i], P_ROW[_k]))
for _k, _f2 in M_C.items():
    CELDA[_k] = (HC, 'B%d' % _f2)
for _k, _f2 in M_A.items():
    CELDA[_k] = (HP, 'B%d' % _f2)
for _k, _f2 in E.items():
    CELDA[_k] = (HE, 'B%d' % _f2)
for _j in range(NPR):
    for _k, _c in (('bo', CBO), ('lq', CLQ), ('pn', CPN), ('ot', COT), ('ra', CRA), ('pvp', CPVP), ('ud', CUD)):
        CELDA['%s%d' % (_k, _j)] = (HE, '%s%d' % (_c, E_P0 + _j))


def celdas_salida(r):
    """(hoja, celda, valor esperado, etiqueta) de todas las salidas."""
    out = []
    for i, c in enumerate(r['Pe']):
        for k in S_C:
            out.append((HC, '%s%d' % (PC[i], S_C[k]), c[k], 'Pe%d.%s' % (i, k)))
    for k in S_D:
        out.append((HC, 'B%d' % S_D[k], r['D'][k], 'D.' + k))
    for k in S_A:
        out.append((HP, 'B%d' % S_A[k], r['A'][k], 'A.' + k))
    out.append((HE, 'B%d' % E_BO, r['bola'], 'bola'))
    out.append((HE, 'B%d' % E_PA, r['pata'], 'pata'))
    for j, q in enumerate(r['Pr']):
        for k, col, _fm in SALP:
            out.append((HE, '%s%d' % (col, E_P0 + j), q[k], '%s[%d]' % (k, j)))
    for k in T:
        out.append((HE, 'B%d' % T[k], r['T'][k], 'T.' + k))
    return out


def _distinto(x, v):
    if isinstance(v, str):
        return x != v
    return not isinstance(x, (int, float)) or abs(float(x) - v) > 0.01


def _checks_logica(nombre, jj, err):
    r = calcular(jj)
    for i, c in enumerate(r['Pe']):
        if abs(c['cku'] * jj['ut%d' % i] * (1 - jj['m']) - jj['pr%d' % i]) > 1e-6:
            err.append('%s: coste útil × útil × (1 − merma) ≠ precio en la columna %d' % (nombre, i))
        if abs(c['cku_mo'] * jj['ut%d' % i] * (1 - jj['m']) - (jj['pr%d' % i] + jj['tl%d' % i] / 60.0 * jj['hh'])) > 1e-6:
            err.append('%s: coste útil con mano de obra no cuadra en la columna %d' % (nombre, i))
        if abs(c['c_bola'] * jj['KG'] / jj['gb'] - c['cku_mo']) > 1e-6:
            err.append('%s: coste de bola ≠ coste útil × gramos ÷ 1.000 en la columna %d' % (nombre, i))
        if abs(c['c_dos'] - 2 * c['c_bola']) > 1e-9 or abs(c['c_tres'] - 3 * c['c_bola']) > 1e-9:
            err.append('%s: coste de 2 y 3 bolas ≠ 2 y 3 × bola en la columna %d' % (nombre, i))
    d = r['D']
    if abs(d['bkg'] * jj['gb'] - 1000) > 1e-6 or abs(d['bkn'] - d['bkg'] * (1 - jj['m'])) > 1e-6:
        err.append('%s: bolas por kg ≠ 1.000 ÷ g (o con la merma)' % nombre)
    if abs(d['carne_dia'] * 1000 - jj['cd'] * jj['bc'] * jj['gb']) > 1e-6:
        err.append('%s: kg de carne al día ≠ comensales × bolas × g ÷ 1.000' % nombre)
    a = r['A']
    if abs(a['pat_cong'] * 1000 - jj['gc'] * jj['pc']) > 1e-6:
        err.append('%s: patata congelada por ración ≠ g × €/kg ÷ 1.000' % nombre)
    if abs(a['kg_bruto'] * jj['uf'] * 1000 - jj['gf']) > 1e-6:
        err.append('%s: kilos sin pelar × aprovechamiento ≠ gramos cortados' % nombre)
    if abs(a['rac_dia'] - jj['cd'] * jj['rpc']) > 1e-9:
        err.append('%s: raciones al día ≠ comensales × raciones por comensal' % nombre)
    if abs(a['c_cong'] - a['pat_cong'] - a['ace_cong']) > 1e-9:
        err.append('%s: ración congelada ≠ patata + aceite' % nombre)
    for j, q in enumerate(r['Pr']):
        if not q['plat'] < q['llevar'] < q['sala']:
            err.append('%s: margen plataforma < para llevar < sala no se cumple en el producto %d' % (nombre, j))
    q = r['Pr'][J_MENU]
    if not q['plat'] < q['llevar'] < q['sala']:
        err.append('%s: margen del menú plataforma < para llevar < sala no se cumple' % nombre)
    if abs(r['D']['X'] - (r['Pe'][0]['cku_mo'] - r['Pe'][1]['cku'])) > 1e-9:
        err.append('%s: X ≠ coste útil del blend − coste útil de la pieza' % nombre)


def autotest(destino):
    from pycel import ExcelCompiler
    err, n = [], 0
    wv = openpyxl.load_workbook(destino, data_only=True)
    wf = openpyxl.load_workbook(destino)
    base = valores_defecto()
    j2 = dict(base, pr0=9.5, ut0=0.98, pr1=8.0, ut1=0.85, tl1=9.0, pr2=11.0, pr3=10.0, ut3=0.9, tl3=2.0, m=0.05, gb=80.0, mc=0.3, hh=18.0, bc=2.0,
              cd=120, cong=OPC[1][0], pc=1.8, gc=170.0, ac=0.06, pf=0.8, gf=200.0, uf=0.7, af=0.12, tf=8.0, hp=18.0, lc=12.0, nc=3, pa=2.2,
              rcm=700, rpc=0.9, dm=26, iva=0.10, obj=0.30, tol=0.04, sel=2, tp=2, pan=0.6, qk=9.0, gq=25.0, cl=0.5, cp=0.9, rec=0.1, com=0.25,
              bo3=3, lq2=1, ra8=2, pvp1=12.5, ud0=100, ud12=300)
    j3 = dict(base, pr0=8.0, ut0=0.97, pr1=6.9, ut1=0.92, tl1=4.0, pr2=12.5, pr3=7.0, m=0.0, gb=120.0, mc=0.2, hh=14.0, bc=1.2, cd=60,
              cong=OPC[2][0], pc=2.1, gc=130.0, ac=0.04, pf=0.5, gf=180.0, uf=0.8, af=0.08, tf=4.0, hp=14.0, lc=8.0, nc=1, pa=1.5, rcm=1200,
              rpc=0.7, dm=22, iva=0.10, obj=0.35, tol=0.1, sel=3, tp=1, pan=0.4, qk=7.0, gq=15.0, cl=0.3, cp=0.8, rec=0.2, com=0.35,
              bo0=2, lq1=3, pn4=2, pvp9=7.5, pvp12=17.0, ud5=0, ud12=200)
    # 1. caché y pycel del juego por defecto; pycel de los otros dos con set_value
    xl = ExcelCompiler(filename=destino)
    r0 = calcular(base)
    for hoja, c, v, lab in celdas_salida(r0):
        n += 1
        for et, x in (('pycel', xl.evaluate("'%s'!%s" % (hoja, c))), ('caché', wv[hoja][c].value)):
            if _distinto(x, v):
                err.append('defecto · %s!%s (%s): %s %r ≠ esperado %r' % (hoja, c, lab, et, x, v))
    for nombre, jj in (('juego 2', j2), ('juego 3', j3)):
        r = calcular(jj)
        xl2 = ExcelCompiler(filename=destino)
        for hoja, c, v, lab in celdas_salida(r0):
            xl2.evaluate("'%s'!%s" % (hoja, c))
        for k, v in jj.items():
            if v != base[k]:
                h, c = CELDA[k]
                xl2.set_value("'%s'!%s" % (h, c), v)
        for hoja, c, v, lab in celdas_salida(r):
            n += 1
            x = xl2.evaluate("'%s'!%s" % (hoja, c))
            if _distinto(x, v):
                err.append('%s · %s!%s (%s): pycel %r ≠ esperado %r' % (nombre, hoja, c, lab, x, v))
    # 2. identidades de la lógica en los tres juegos
    for nombre, jj in (('defecto', base), ('juego 2', j2), ('juego 3', j3)):
        n += 20
        _checks_logica(nombre, jj, err)
    # identidades sobre la hoja viva (pycel, valores por defecto): coste útil = precio ÷ útil ÷ (1 − merma), bolas por kg = 1.000 ÷ g
    ev = lambda h, c: xl.evaluate("'%s'!%s" % (h, c))
    for i in range(NPE):
        n += 1
        pr_, ut_ = ev(HC, '%s%d' % (PC[i], P_ROW['pr'])), ev(HC, '%s%d' % (PC[i], P_ROW['ut']))
        m_ = ev(HC, 'B%d' % M_C['m'])
        if abs(ev(HC, '%s%d' % (PC[i], S_C['cku'])) - pr_ / ut_ / (1 - m_)) > 1e-6:
            err.append('hoja Carne, columna %s: coste útil ≠ precio ÷ útil ÷ (1 − m)' % PC[i])
    gb_ = ev(HC, 'B%d' % M_C['gb'])
    n += 2
    if abs(ev(HC, 'B%d' % S_D['bkg']) - 1000.0 / gb_) > 1e-9:
        err.append('hoja Carne: bolas por kg ≠ 1.000 ÷ g')
    if abs(ev(HC, 'B%d' % S_D['bkn']) - 1000.0 * (1 - ev(HC, 'B%d' % M_C['m'])) / gb_) > 1e-9:
        err.append('hoja Carne: bolas por kg con la merma ≠ 1.000 × (1 − m) ÷ g')
    # 3. valores con sentido
    sim, dob, men = r0['Pr'][J_SIM], r0['Pr'][J_DOB], r0['Pr'][J_MENU]
    if not 0.18 <= sim['fc'] <= 0.26:
        err.append('food cost de la simple %.1f %% fuera de 18-26 %%' % (sim['fc'] * 100))
    if not 0.24 <= dob['fc'] <= 0.30:
        err.append('food cost de la doble %.1f %% fuera de 24-30 %%' % (dob['fc'] * 100))
    if not 0.22 <= r0['T']['fcp'] <= 0.27:
        err.append('food cost ponderado por defecto %.1f %% fuera de 22-27 %%' % (r0['T']['fcp'] * 100))
    if not men['plat'] < men['llevar'] < men['sala']:
        err.append('margen del menú: plataforma < para llevar < sala no se cumple')
    ck = (r0['Pe'][0]['cku'], r0['Pe'][1]['cku'])
    if abs(ck[0] - 9.18) > 0.01 or abs(ck[1] - 8.59) > 0.01:
        err.append('coste por kg útil por defecto %r ≠ 9,18 · 8,59' % [round(x, 3) for x in ck])
    D_ = r0['D']
    if abs(D_['bkg'] - 10) > 1e-9 or abs(D_['bkn'] - 9.7) > 1e-9 or abs(D_['carne_dia'] - 15.2) > 0.01:
        err.append('bolas por defecto: 10 · 9,7 · 15,2 kg al día (salió %r)' % [round(D_[k], 3) for k in ('bkg', 'bkn', 'carne_dia')])
    A_ = r0['A']
    if abs(A_['c_cong'] - 0.28) > 0.005 or abs(A_['c_fre'] - 0.19) > 0.01 or abs(A_['rac_dia'] - 76) > 1e-9:
        err.append('patatas por defecto: 0,28 · 0,19 · 76 (salió %r)' % [round(A_[k], 3) for k in ('c_cong', 'c_fre', 'rac_dia')])
    if abs(dob['c_tot'] - 2.90) > 0.02 or abs(sim['c_tot'] - 1.80) > 0.02 or abs(r0['Pr'][2]['c_tot'] - 3.99) > 0.02 or abs(r0['Pr'][3]['c_tot'] - 3.16) > 0.02:
        err.append('escandallo por defecto: simple 1,80 · doble 2,90 · triple 3,99 · bacon 3,16 (salió %r)' % [round(r0['Pr'][k]['c_tot'], 3) for k in range(4)])
    if abs(men['sala'] - 10.83) > 0.02 or abs(men['llevar'] - 10.38) > 0.02 or abs(men['plat'] - 7.43) > 0.05:
        err.append('margen del menú por defecto: 10,83 · 10,38 · 7,43 (salió %r)' % [round(men[k], 3) for k in ('sala', 'llevar', 'plat')])
    if D_['sem'] != SEM_BLEND or A_['sem'] != SEM_CONG:
        err.append('veredictos por defecto: «%s» y «%s» (esperado blend y congelada)' % (D_['sem'], A_['sem']))
    if calcular(dict(base, tl1=0.0))['D']['sem'] != SEM_PICAR:
        err.append('el veredicto no pasa a «picar en el local» con 0 minutos de limpieza')
    if calcular(dict(base, tf=0.0))['A']['sem'] != SEM_FRE:
        err.append('el veredicto no pasa a «la fresca compensa» con 0 minutos de pelado')
    if sorted(q['rank'] for q in r0['Pr'])[0] != 1:
        err.append('los puestos por margen no son coherentes')
    # 4. valores expuestos: caché frente a calcular()
    esperados = {'COSTE_KG_UTIL_BLEND': r0['Pe'][0]['cku'], 'COSTE_KG_UTIL_PIEZA': r0['Pe'][1]['cku'], 'BOLAS_KG': D_['bkg'],
                 'BOLAS_KG_NETAS': D_['bkn'], 'CARNE_DIA_KG': D_['carne_dia'], 'COSTE_SIMPLE': sim['c_tot'], 'COSTE_DOBLE': dob['c_tot'],
                 'FC_SIMPLE': sim['fc'], 'FC_DOBLE': dob['fc'], 'FC_PONDERADO': r0['T']['fcp'], 'COSTE_RACION_PATATA_CONG': A_['c_cong'],
                 'COSTE_RACION_PATATA_FRESCA': A_['c_fre'], 'ACEITE_RACION': A_['ace_cong'], 'MARGEN_SALA_MENU': men['sala'],
                 'MARGEN_PLATAFORMA_MENU': men['plat']}
    expuestos = valores_expuestos(destino)
    if list(expuestos) != list(esperados):
        err.append('EXPUESTOS ≠ lista de F1-camino §4: %r' % list(expuestos))
    for nombre, e_ in esperados.items():
        n += 1
        if _distinto(expuestos[nombre], e_):
            err.append('expuesto %s: caché %r ≠ %r' % (nombre, expuestos[nombre], e_))
    # 5. formato, protección, DV, fórmulas, textos
    n += 1
    fg = motor.gate_formatos(wf, FICHERO)
    if fg:
        err.append('gate_formatos: %r' % fg[:3])
    dvh = dict((w.title, dict((str(rng), dv) for dv in w.data_validations.dataValidation for rng in dv.sqref.ranges))
               for w in (wf[HC], wf[HP], wf[HE]))
    for hoja, k, tipo in ((HC, 'pr0', 'decimal'), (HC, 'ut1', 'decimal'), (HC, 'm', 'decimal'), (HC, 'cong', 'list'), (HC, 'cd', 'whole'),
                          (HP, 'pc', 'decimal'), (HP, 'nc', 'whole'), (HE, 'obj', 'decimal'), (HE, 'iva', 'decimal'), (HE, 'com', 'decimal'),
                          (HE, 'sel', 'whole'), (HE, 'tp', 'whole'), (HE, 'bo0', 'decimal'), (HE, 'ud0', 'decimal')):
        n += 1
        c = CELDA[k][1]
        dv = dvh[hoja].get(c)
        if dv is None or dv.type != tipo:
            err.append('DV ausente o de otro tipo en %s!%s (esperada %s)' % (hoja, c, tipo))
    textos, todo_celdas = [], []
    for w in wf.worksheets:
        n += 1
        if w.page_setup.paperSize != 9 or not w.protection.sheet:
            err.append('%s: A4 y protección' % w.title)
        for row in w.iter_rows():
            for c in row:
                v = c.value
                if v is not None:
                    todo_celdas.append(str(v))
                if isinstance(v, str) and c.data_type != 'f':
                    textos.append(v)
                if c.data_type == 'f':
                    if wv[w.title][c.coordinate].value is None:
                        err.append('%s!%s: fórmula sin caché' % (w.title, c.coordinate))
                    if len(str(c.value)) >= 200:
                        err.append('%s!%s: fórmula de %d caracteres (≥ 200)' % (w.title, c.coordinate, len(str(c.value))))
                    sin_refs = re.sub(r'"[^"]*"', '', re.sub(r"'[^']+'!|\$?[A-Z]{1,2}\$?\d+", '', str(c.value)))
                    for num in re.findall(r'(?<![A-Za-z])\d+\.?\d*', sin_refs):
                        if num not in ('0', '1'):
                            err.append('%s!%s: número tecleado %s en %s' % (w.title, c.coordinate, num, c.value))
                if c.fill.fgColor.rgb in ('00E8F5E9', 'FFE8F5E9') and c.protection.locked:
                    err.append('%s!%s: verde bloqueada' % (w.title, c.coordinate))
                if c.fill.fgColor.rgb in ('00E8F5E9', 'FFE8F5E9') and c.data_type == 'f':
                    err.append('%s!%s: verde con fórmula' % (w.title, c.coordinate))
    for hoja, ks, mm in ((HC, ('KG', 'MH'), M_C), (HP, ('KG', 'ML', 'MH'), M_A)):
        for k in ks:
            n += 1
            c = wf[hoja]['B%d' % mm[k]]
            rot = wf[hoja]['A%d' % mm[k]].value
            if c.fill.fgColor.rgb not in ('00F2F2F2', 'FFF2F2F2') or c.protection.locked is not True or 'constante' not in rot:
                err.append('constante gris %s mal marcada o desbloqueada' % k)
    todo = '\n'.join(textos)
    barrido = '\n'.join(todo_celdas)
    for rx in PROHIBIDOS:
        if re.search(rx, barrido, re.I):
            err.append('prohibido /%s/ en %r' % (rx, re.search(rx, barrido, re.I).group(0)))
    for m_ in re.finditer(r'obligatori', barrido, re.I):                           # cerca de la cocción o de los 70 °C
        ctx = barrido[max(0, m_.start() - 80):m_.end() + 80]
        if re.search(r'cocci[oó]n|70 ?°C', ctx, re.I):
            err.append('«obligatori» cerca de cocción o 70 °C: %r' % ctx)
    for m_ in re.finditer(r'7 días', barrido):
        if re.search(r'congel|picad|bola|−18', barrido[max(0, m_.start() - 80):m_.end() + 80], re.I):
            err.append('«7 días» cerca de congelar, picar o bola')
    for m_ in re.finditer(r'art\. 30', barrido):
        if 'RD 1086/2020' not in barrido[max(0, m_.start() - 60):m_.end() + 60]:
            err.append('«art. 30» sin «RD 1086/2020» al lado')
    for m_ in re.finditer(r'de un día para otro', barrido):
        if not re.search(r'(nunca|no|ni|jamás)\s+(se\s+)?(guarda\w*\s+)?(de\s+)?$', barrido[max(0, m_.start() - 30):m_.start()]):
            err.append('«de un día para otro» sin negación delante: %r' % barrido[max(0, m_.start() - 40):m_.end()])
    avisos = [rx for rx in AVISOS if re.search(rx, barrido, re.I)]
    if re.search(r'[^\x00-ɏ‐-‧€−→·≥≈]', barrido):
        err.append('caracteres no latinos o raros: %r' % sorted(set(re.findall(r'[^\x00-ɏ‐-‧€−→·≥≈]', barrido))))
    if 'estimación del kit, editable' not in barrido.lower():
        err.append('falta el rótulo «estimación del kit, editable»')
    if 'sin fuente: mídelo' not in barrido.lower():
        err.append('falta «sin fuente: mídelo»')
    if re.search(r'\bunidades\b', barrido, re.I):
        err.append('aparece «unidades» (los rótulos son recuentos de bolas, raciones y productos)')
    for rot in ('Hamburguesa doble (2 bolas)', 'Ración de patatas'):
        n += 1
        if rot not in barrido:
            err.append('falta el rótulo «%s»' % rot)
    if motor.VERSION_FMT.format(pid=PID) not in todo or motor.BIO_LINE not in todo:
        err.append('falta la línea de versión o la bio')
    if '1.0' not in motor.VERSION_FMT or 'noviembre 2026' not in motor.VERSION_FMT:
        err.append('versión no es 1.0 · noviembre 2026: %s' % motor.VERSION_FMT)
    if wf.sheetnames != [HC, HP, HE, HI]:
        err.append('pestañas %r' % wf.sheetnames)
    if NPR != 14:
        err.append('no son 14 productos')
    p = wf.properties
    if (p.creator != 'AI Chef Pro' or p.description != 'aichef.pro/%s' % PID or 'AI Chef Pro' not in (p.keywords or '')
            or p.subject != 'Plan de Negocio: Hamburguesería Smash'):
        err.append('docProps: creator=%r description=%r keywords=%r subject=%r' % (p.creator, p.description, p.keywords, p.subject))
    for x in err:
        print('ERROR', x)
    for x in avisos:
        print('AVISO /%s/ presente: revisar que sea solo como prohibición o con su salvedad' % x)
    print('autotest calculadora hamburguesería ES: %d comprobaciones (pycel + caché por defecto; pycel en 2 juegos más; identidades) · %s'
          % (n, 'VERDE' if not err else 'ROJO (%d errores)' % len(err)))
    print('  carne (€/kg útil): ' + ' · '.join('%s %.3f' % (CARNES[i][0], c['cku']) for i, c in enumerate(r0['Pe'])))
    print('  picar: mano de obra %.3f €/kg útil · compensa hasta %.3f €/kg útil · %s' % (D_['lab_kg'], D_['X'], D_['sem']))
    print('  bolas: %.2f por kg en crudo · %.2f con la merma · %.1f bolas/día · %.2f kg de carne/día' % (D_['bkg'], D_['bkn'], D_['bdia'], D_['carne_dia']))
    print('  patatas: congelada %.4f € (aceite %.4f) · fresca %.4f € sin MO, %.4f con MO · la fresca compensa hasta %.3f €/kg (MO %.3f) · %s'
          % (A_['c_cong'], A_['ace_cong'], A_['c_fre'], A_['c_fre_mo'], A_['X'], A_['mo_kg'], A_['sem']))
    print('  patatas: %.0f raciones/día · %.1f kg congelada/día · %.1f kg fresca/día · %.1f l de aceite/mes' % (
        A_['rac_dia'], A_['kg_cong_dia'], A_['kg_fre_dia'], A_['l_mes']))
    for nombre, (h, c) in EXPUESTOS.items():
        print('  %-28s %s!%-5s %s' % (nombre, h, c, expuestos[nombre]))
    for j, p_ in enumerate(PRODS):
        q = r0['Pr'][j]
        print('  %-42s coste %.3f € · food cost %5.1f %% · %s · margen sala %.2f / llevar %.2f / plataforma %.2f € · puesto %d · %d/semana'
              % (p_[0], q['c_tot'], q['fc'] * 100, q['sem'], q['sala'], q['llevar'], q['plat'], q['rank'], p_[7]))
    print('  ponderado: ventas %.2f € · coste %.2f € · food cost %.2f %% · %s' % (r0['T']['ing'], r0['T']['cos'], r0['T']['fcp'] * 100, r0['T']['semp']))
    print('  menú: caída del margen en plataforma %.1f %%' % ((men['plat'] / men['sala'] - 1) * 100))
    return 1 if err else 0


if __name__ == '__main__':
    ap = argparse.ArgumentParser()
    ap.add_argument('--salida', required=True, help='carpeta de salida')
    ap.add_argument('--autotest', action='store_true')
    a = ap.parse_args()
    destino = os.path.join(a.salida, FICHERO)
    if not a.autotest and generar(destino):
        sys.exit('ERROR: fallos de pycel en inject_cache')
    sys.exit(autotest(destino))
