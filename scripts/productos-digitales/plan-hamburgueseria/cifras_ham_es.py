#!/usr/bin/env python3
"""
cifras_ham_es.py — Plan de Negocio: Hamburguesería Smash ES 1.0 (copia de `plan-sushi/cifras_sus_es.py`, la misma
rejilla: comensal y rotación encendida): lee la CACHÉ del plan financiero (después de `inject_cache`, que hace `main.py`)
por RÓTULO, escribe `cifras_caso_es.json` (los tokens {ASI} del docx) y pasa los gates del caso (D9 = Sushi Bar;
F1-camino §2 y §7 F2): DSCR mínimo ≥ 1,25 · saldo mínimo de tesorería > 0 · cobertura de horas ≥ 98 % · holgura sobre
el equilibrio de caja del año 1 ≥ 15 % (las de los años 2 y 3, informativas, ≥ 0) · exceso de financiación (origen −
usos) ≥ 0 y < 5 % (RD-34) · resultado neto a régimen (años 2 y 3) entre el 7 y el 12 % de las ventas, y el del año 1
≤ 12 % · rotación de equilibrio de caja ≤ techo del local · `pct_delivery` < 25 % (demo §2.11.9 del motor, que prueba
con 25 %) · `PyG!G` de la línea de bebida = H × 21 % + (1 − H) × 10 %, con la fórmula en el libro · Mes 1 = enero
(`Supuestos!B60/B61` = 7/12, `Tesorería!B5:M5` = la tupla de `a.ESTACIONALIDAD`, mes del saldo mínimo coherente) ·
`prueba_hamburgueseria` (F1-camino §3 y §7 F2 (3)-(4): techo de rotación 3,5 y caso 2,5-2,8; aforo 36; plataforma
< 25 %; 0 restos de la semilla y de sus moldes en el plan; rótulos «Comensales/día»; payback como número).

Tokens nuevos frente a la semilla (F1-camino §2 (2)): `PEDIDOS_DIA` (≈ 25: los equivalentes de llevar y domicilio ÷
1,7), `PEDIDOS_LLEVAR_DIA` y `PEDIDOS_PLATAFORMA_DIA` (13 y 12, F1-camino §2), `COMENSALES_SALA` (≈ 53), `ROT_SALA`
(≈ 1,5), `VENTAS_21_PCT` (la columna H de la bebida sobre las ventas: 0,16 × 0,28 ≈ 4,5 %) y las partidas de la
hamburguesería (`PLANCHA`, `FREIDORAS`, `CAMPANA`, `MESA_REFRIGERADA`, `FRIO`, `BARRA`, `MANTENIMIENTO_ANO`).

Con `--calc` lee TODO el `EXPUESTOS` de la calculadora de la carne, las patatas y el escandallo (`calculadora_ham_es.py`):
coste por kg útil del blend y de la pieza, bolas por kg (en crudo y con la merma de cierre), kg de carne al día, coste y
food cost de la smash simple y la doble, food cost ponderado, coste de la ración de patatas congelada y fresca, aceite por
ración y el margen del menú en sala y en plataforma. No entra en los gates y no falla si la calculadora aún no existe.
Corre en el VPS (openpyxl).

    /root/venv-guias/bin/python cifras_ham_es.py --xlsx <plan-financiero-hamburgueseria.xlsx> [--calc <calculadora…xlsx>]

Sesión Claude Code, 31-oct-2026.
"""
import argparse
import importlib
import json
import math
from decimal import ROUND_HALF_UP, Decimal
import os
import re
import sys

import openpyxl

AQUI = os.path.dirname(os.path.abspath(__file__))
SALIDA = os.path.join(AQUI, 'cifras_caso_es.json')
sys.path.insert(0, AQUI)
sys.path.insert(0, os.path.join(os.path.dirname(AQUI), 'planes-v2_0'))

#: Mes 1 = enero (F1-camino §2): nombre del mes del calendario de cada «Mes n» de la Tesorería.
MESES = ('enero', 'febrero', 'marzo', 'abril', 'mayo', 'junio', 'julio', 'agosto', 'septiembre', 'octubre',
         'noviembre', 'diciembre')

#: token → (hoja, regex del rótulo de la columna A, columna, formato). Formatos: eur (miles con punto, sin
#: decimales) · eur2 · pct (1 decimal) · pct0 · dec1 · dec2 · int · int_up (redondeo hacia arriba).
SUP, PYG, INV, PE, TES, FIN, PER, ESC = ('0. Supuestos', 'PyG 3 Años', 'Inversión Inicial', 'Punto Equilibrio',
                                         'Tesorería 12 meses', 'Financiación', 'Personal', 'Escenarios')
MAPA = {
    'COMENSALES_DIA': (SUP, r'^comensales/d[ií]a \(media', 2, 'int'),
    'TICKET': (SUP, r'^ticket medio sin iva', 2, 'eur2'),
    'TICKET_IVA': (SUP, r'^pvp equivalente con iva', 2, 'eur2'),
    'DIAS_APERTURA': (SUP, r'^d[ií]as de apertura al a[ñn]o', 2, 'int'),
    'CREC_A2': (SUP, r'^crecimiento de volumen del a[ñn]o 2', 2, 'pct0'),
    'CREC_A3': (SUP, r'^crecimiento de volumen del a[ñn]o 3', 2, 'pct0'),
    'PCT_COMIDA': (SUP, r'^ventas de comida sobre el total', 2, 'pct0'),
    'PCT_BEBIDA': (SUP, r'^ventas de bebida sobre el total', 2, 'pct0'),
    'FOOD_COST': (SUP, r'^coste de mercanc[ií]a sobre las ventas de comida', 2, 'pct'),
    'BEBIDA_COST': (SUP, r'^coste de mercanc[ií]a sobre las ventas de bebida', 2, 'pct0'),
    'CONSUMIBLES_PCT': (SUP, r'^consumibles sobre ventas', 2, 'pct'),
    'DELIVERY_PCT': (SUP, r'^ventas por delivery', 2, 'pct0'),
    'COMISION_DELIVERY': (SUP, r'^comisi[oó]n de la plataforma', 2, 'pct0'),
    'COMISION_TPV': (SUP, r'^comisi[oó]n de los medios de pago', 2, 'pct'),
    'ALQUILER_MES': (SUP, r'^alquiler mensual del local', 2, 'eur'),
    'FIANZA_MESES': (SUP, r'^fianza del alquiler', 2, 'int'),
    'SUMINISTROS_MES': (SUP, r'^suministros mensuales', 2, 'eur'),
    'SEGUROS_ANO': (SUP, r'^seguros', 2, 'eur'),
    'RECURSOS_PROPIOS': (SUP, r'^recursos propios', 2, 'eur'),
    'TIPO_PRESTAMO': (SUP, r'^tipo de inter[eé]s nominal', 2, 'pct0'),
    'PLAZO_PRESTAMO': (SUP, r'^plazo del pr[eé]stamo', 2, 'int'),
    'CARENCIA': (SUP, r'^carencia de principal', 2, 'int'),
    'MESES_FONDO': (SUP, r'^fondo de maniobra \(meses', 2, 'int'),
    'IS_NUEVA': (SUP, r'^impuesto de sociedades, nueva creaci', 2, 'pct0'),
    'IS_GENERAL': (SUP, r'^impuesto de sociedades, tipo general', 2, 'pct0'),
    'VIDA_OBRA': (SUP, r'^vida [uú]til de obra', 2, 'int'),
    'VIDA_MAQUINARIA': (SUP, r'^vida [uú]til de maquinaria', 2, 'int'),
    'ROTACION_IMPLICITA': (SUP, r'^rotaciones al d[ií]a impl[ií]citas', 2, 'dec1'),
    'AFORO': (SUP, r'^aforo del local', 2, 'int'),
    'RAMPA_MES1': (SUP, r'^actividad del mes 1', 2, 'pct0'),
    'MESES_RAMPA': (SUP, r'^meses hasta alcanzar', 2, 'int'),
    'MESES_RENTA_PREVIA': (SUP, r'^meses de alquiler antes de abrir', 2, 'int'),
    'DIAS_COBRO': (SUP, r'^d[ií]as medios de cobro', 2, 'int'),
    'DIAS_PAGO': (SUP, r'^d[ií]as medios de pago', 2, 'int'),
    'MES_PAGA_1': (SUP, r'^mes de la paga extra de verano', 2, 'int'),
    'MES_PAGA_2': (SUP, r'^mes de la paga extra de navidad', 2, 'int'),
    'BEBIDA_COMPRA_21': (SUP, r'^bebida que compras al 21', 2, 'pct0'),
    'HOLGURA_MINIMA': (SUP, r'^holgura m[ií]nima sobre el punto de equilibrio', 2, 'pct0'),
    'VENTAS_A1': (PYG, r'^ingresos totales', 2, 'eur'),
    'VENTAS_A2': (PYG, r'^ingresos totales', 3, 'eur'),
    'VENTAS_A3': (PYG, r'^ingresos totales', 4, 'eur'),
    'VENTAS_COMIDA_A1': (PYG, r'^ventas de hamburguesas', 2, 'eur'),
    'VENTAS_BEBIDA_A1': (PYG, r'^ventas de bebida', 2, 'eur'),
    'IVA_BEBIDA_LINEA': (PYG, r'^ventas de bebida', 7, 'pct'),
    'BEBIDA_21_PCT': (PYG, r'^ventas de bebida', 8, 'pct'),          # 28,0 % (F1-camino §2 (2): en `pct`)
    'IVA_COMPRA_COMIDA': (PYG, r'^coste de mercanc[ií]a — comida', 7, 'pct'),
    'IVA_COMPRA_BEBIDA': (PYG, r'^coste de mercanc[ií]a — bebida', 7, 'pct'),
    'COMENSALES_DIA_A2': (PYG, r'^comensales/d[ií]a \(media\)', 3, 'int'),
    'COMENSALES_DIA_A3': (PYG, r'^comensales/d[ií]a \(media\)', 4, 'int'),
    'MERCANCIA_COMIDA_A1': (PYG, r'^coste de mercanc[ií]a — comida', 2, 'eur'),
    'MERCANCIA_BEBIDA_A1': (PYG, r'^coste de mercanc[ií]a — bebida', 2, 'eur'),
    'CONSUMIBLES_A1': (PYG, r'^consumibles y envases', 2, 'eur'),
    'COMISIONES_DELIVERY_A1': (PYG, r'^comisiones de delivery', 2, 'eur'),
    'COSTES_VARIABLES_A1': (PYG, r'^total costes variables', 2, 'eur'),
    'MARGEN_BRUTO_A1': (PYG, r'^margen bruto$', 2, 'eur'),
    'MARGEN_BRUTO_PCT_A1': (PYG, r'^margen bruto / ventas', 2, 'pct'),
    'ALQUILER_ANO': (PYG, r'^alquiler del local', 2, 'eur'),
    'ALQUILER_PCT_A1': (PYG, r'^alquiler / ventas', 2, 'pct'),
    'PERSONAL_A1': (PYG, r'^personal \(n[oó]minas', 2, 'eur'),
    'PERSONAL_A2': (PYG, r'^personal \(n[oó]minas', 3, 'eur'),
    'PERSONAL_A3': (PYG, r'^personal \(n[oó]minas', 4, 'eur'),
    'PERSONAL_PCT_A1': (PYG, r'^coste de personal / ventas', 2, 'pct'),
    'MERCANCIA_PCT_A1': (PYG, r'^coste de mercanc[ií]a / ventas', 2, 'pct'),
    'SUMINISTROS_ANO': (PYG, r'^suministros', 2, 'eur'),
    'GESTORIA_ANO': (PYG, r'^gestor[ií]a', 2, 'eur'),
    'MANTENIMIENTO_ANO': (PYG, r'^mantenimiento de la campana', 2, 'eur'),
    'MARKETING_ANO': (PYG, r'^marketing y publicidad', 2, 'eur'),
    'RESIDUOS_ANO': (PYG, r'^gesti[oó]n de residuos', 2, 'eur'),
    'PRL_FORMACION_ANO': (PYG, r'^prevenci[oó]n de riesgos laborales', 2, 'eur'),
    'AMORTIZACION_A1': (PYG, r'^amortizaci[oó]n del inmovilizado', 2, 'eur'),
    'INTERESES_A1': (PYG, r'^gastos financieros', 2, 'eur'),
    'FIJOS_A1': (PYG, r'^total costes fijos', 2, 'eur'),
    'AMORTIZACION_A2': (PYG, r'^amortizaci[oó]n del inmovilizado', 3, 'eur'),
    'AMORTIZACION_A3': (PYG, r'^amortizaci[oó]n del inmovilizado', 4, 'eur'),
    'INTERESES_A2': (PYG, r'^gastos financieros', 3, 'eur'),
    'INTERESES_A3': (PYG, r'^gastos financieros', 4, 'eur'),
    'RAI_A1': (PYG, r'^resultado antes de impuestos', 2, 'eur'),
    'RAI_A2': (PYG, r'^resultado antes de impuestos', 3, 'eur'),
    'RAI_A3': (PYG, r'^resultado antes de impuestos', 4, 'eur'),
    'IS_A1': (PYG, r'^impuesto de sociedades$', 2, 'eur'),
    'IS_A3': (PYG, r'^impuesto de sociedades$', 4, 'eur'),
    'NETO_A1': (PYG, r'^resultado neto$', 2, 'eur'),
    'NETO_A2': (PYG, r'^resultado neto$', 3, 'eur'),
    'NETO_A3': (PYG, r'^resultado neto$', 4, 'eur'),
    'NETO_PCT_A1': (PYG, r'^resultado neto / ventas', 2, 'pct'),
    'NETO_PCT_A2': (PYG, r'^resultado neto / ventas', 3, 'pct'),
    'NETO_PCT_A3': (PYG, r'^resultado neto / ventas', 4, 'pct'),
    'CAPEX': (INV, r'^capex', 2, 'eur'),
    'FONDO_MANIOBRA': (INV, r'^fondo de maniobra$', 2, 'eur'),
    'INVERSION_TOTAL': (INV, r'^inversi[oó]n total', 2, 'eur'),
    'IVA_INVERSION': (INV, r'^iva soportado sobre la inversi', 2, 'eur'),
    'NECESIDAD_CAJA': (INV, r'^necesidad total de caja', 2, 'eur'),
    'FIANZA': (INV, r'^fianza del alquiler', 2, 'eur'),
    'RENTA_PREVIA': (INV, r'^alquiler de los meses previos', 2, 'eur'),
    'IMPREVISTOS': (INV, r'^imprevistos de obra', 2, 'eur'),
    'PROYECTO': (INV, r'^proyecto t[eé]cnico', 2, 'eur'),
    'OBRA': (INV, r'^obra del local', 2, 'eur'),
    'CONDUCTO': (INV, r'^conducto', 2, 'eur'),
    'INSTALACION_ELECTRICA': (INV, r'^instalaci[oó]n el[eé]ctrica', 2, 'eur'),
    'FONTANERIA': (INV, r'^fontaner[ií]a', 2, 'eur'),
    'PLANCHA': (INV, r'^plancha', 2, 'eur'),
    'FREIDORAS': (INV, r'^dos freidoras', 2, 'eur'),
    'CAMPANA': (INV, r'^campana', 2, 'eur'),
    'MESA_REFRIGERADA': (INV, r'^mesa refrigerada', 2, 'eur'),
    'FRIO': (INV, r'^fr[ií]o', 2, 'eur'),
    'BARRA': (INV, r'^barra', 2, 'eur'),
    'LAVAVAJILLAS': (INV, r'^lavavajillas', 2, 'eur'),
    'UTILLAJE': (INV, r'^utillaje', 2, 'eur'),
    'MOBILIARIO': (INV, r'^mobiliario de sala', 2, 'eur'),
    'TPV': (INV, r'^tpv', 2, 'eur'),
    'CONSTITUCION': (INV, r'^constituci[oó]n', 2, 'eur'),
    'MARCA_LANZAMIENTO': (INV, r'^marca, web', 2, 'eur'),
    'STOCK_INICIAL': (INV, r'^stock inicial', 2, 'eur'),
    'BASE_OBRA': (INV, r'^obra, instalaciones y acondicionamiento$', 2, 'eur'),
    'BASE_MAQUINARIA': (INV, r'^maquinaria, mobiliario y equipos', 2, 'eur'),
    'PRESTAMO': (FIN, r'^pr[eé]stamo bancario$', 2, 'eur'),
    'ORIGEN_FONDOS': (FIN, r'^total origen de fondos', 2, 'eur'),
    'DIFERENCIA_FONDOS': (FIN, r'^diferencia \(origen - usos\)', 2, 'eur'),
    'DIFERENCIA_FONDOS_PCT': (FIN, r'^diferencia sobre los usos', 2, 'pct'),
    'CUOTA_ANUAL': (FIN, r'^cuota anual durante la amortizaci', 2, 'eur'),
    'DSCR_A2': (FIN, r'^dscr del a[ñn]o 2', 2, 'dec2'),
    'DSCR_A3': (FIN, r'^dscr del a[ñn]o 3', 2, 'dec2'),
    'DSCR_MIN': (FIN, r'^dscr m[ií]nimo de todo el cuadro', 2, 'dec2'),
    'SALDO_MINIMO': (TES, r'^saldo m[ií]nimo del a[ñn]o', 2, 'eur'),
    'MES_SALDO_MINIMO': (TES, r'^mes en el que la caja toca fondo', 2, 'int'),
    'SALDO_MES12': (TES, r'^saldo acumulado de caja', 13, 'eur'),
    'PAYBACK': (TES, r'^payback del proyecto', 2, 'dec1'),
    'PAYBACK_CAPEX': (TES, r'^payback sobre el capex', 2, 'dec1'),
    'COSTE_VAR_COMENSAL': (PE, r'^coste variable por comensal', 2, 'eur2'),
    'MARGEN_COMENSAL': (PE, r'^margen de contribuci[oó]n por comensal', 2, 'eur2'),
    'BE_COMENSALES_DIA': (PE, r'^comensales necesarios al d[ií]a$', 2, 'int_up'),
    'BE_COMENSALES_DIA_CAJA': (PE, r'^comensales necesarios al d[ií]a \(caja\)', 2, 'int_up'),
    'BE_COMENSALES_DIA_CAJA_A2': (PE, r'^comensales necesarios al d[ií]a \(caja\)', 3, 'int_up'),
    'BE_COMENSALES_DIA_CAJA_A3': (PE, r'^comensales necesarios al d[ií]a \(caja\)', 4, 'int_up'),
    'BE_INGRESOS': (PE, r'^ingresos necesarios al a[ñn]o$', 2, 'eur'),
    'BE_INGRESOS_CAJA': (PE, r'^ingresos necesarios al a[ñn]o \(caja\)', 2, 'eur'),
    'BE_INGRESOS_CAJA_A2': (PE, r'^ingresos necesarios al a[ñn]o \(caja\)', 3, 'eur'),
    'BE_INGRESOS_CAJA_A3': (PE, r'^ingresos necesarios al a[ñn]o \(caja\)', 4, 'eur'),
    'HOLGURA_CAJA': (PE, r'^holgura sobre el equilibrio de caja', 2, 'pct'),
    'HOLGURA_CAJA_A2': (PE, r'^holgura sobre el equilibrio de caja', 3, 'pct'),
    'HOLGURA_CAJA_A3': (PE, r'^holgura sobre el equilibrio de caja', 4, 'pct'),
    'ROT_MAX': (PE, r'^rotaciones al d[ií]a que da el local', 2, 'dec1'),
    'ROT_BE_CAJA': (PE, r'^rotaciones al d[ií]a que exige', 2, 'dec2'),
    'ROT_BE_CAJA_A2': (PE, r'^rotaciones al d[ií]a que exige', 3, 'dec2'),
    'ROT_BE_CAJA_A3': (PE, r'^rotaciones al d[ií]a que exige', 4, 'dec2'),
    'INCORPORACIONES_A2': (PER, r'^coste anual de las incorporaciones del a[ñn]o 2', 2, 'eur'),
    'INCORPORACIONES_A3': (PER, r'^coste anual de las incorporaciones del a[ñn]o 3', 2, 'eur'),
    'COBERTURA_HORAS': (PER, r'^cobertura \(contratadas', 2, 'pct0'),
    'PLANTILLA_PERSONAS': (PER, r'^total plantilla', 2, 'int'),
    'SUELDO_TITULAR': (PER, r'^titular', 4, 'eur'),
    'JORNADAS': (PER, r'^total plantilla', 3, 'dec2'),
    'HORAS_DIA': (PER, r'^horas de servicio al d[ií]a', 2, 'dec1'),
    'HORAS_PRODUCCION': (PER, r'^horas de producci[oó]n', 2, 'dec1'),
    'HORAS_NECESARIAS': (PER, r'^horas de servicio a cubrir al a[ñn]o', 2, 'int'),
    'HORAS_CONTRATADAS': (PER, r'^horas contratadas al a[ñn]o', 2, 'int'),
    'ESC_PES_COMENSALES': (ESC, r'^comensales/d[ií]a$', 2, 'int'),
    'ESC_OPT_COMENSALES': (ESC, r'^comensales/d[ií]a$', 4, 'int'),
    'ESC_PES_TICKET': (ESC, r'^ticket medio sin iva', 2, 'eur2'),
    'ESC_OPT_TICKET': (ESC, r'^ticket medio sin iva', 4, 'eur2'),
    'ESC_PES_DIAS': (ESC, r'^d[ií]as de apertura', 2, 'int'),
    'ESC_OPT_DIAS': (ESC, r'^d[ií]as de apertura', 4, 'int'),
    'ESC_PES_VENTAS': (ESC, r'^ingresos anuales', 2, 'eur'),
    'ESC_OPT_VENTAS': (ESC, r'^ingresos anuales', 4, 'eur'),
    'ESC_PES_NETO': (ESC, r'^resultado neto$', 2, 'eur'),
    'ESC_OPT_NETO': (ESC, r'^resultado neto$', 4, 'eur'),
}

#: Calculadora de la carne, las patatas y el escandallo (`calculadora_ham_es.py`): token → (formato, hoja, regex del
#: rótulo, regex de la cabecera o None). Se leen por las celdas de su `EXPUESTOS` (todas, aunque no estén aquí, con el
#: formato de `_fmt_calc`); hoja None = sólo por `EXPUESTOS`. Los de F1-camino §4.
CALC = {
    'COSTE_KG_UTIL_BLEND': ('eur2', None, None, None),
    'COSTE_KG_UTIL_PIEZA': ('eur2', None, None, None),
    'BOLAS_KG': ('int', None, None, None),
    'BOLAS_KG_NETAS': ('dec1', None, None, None),
    'CARNE_DIA_KG': ('dec1', None, None, None),
    'COSTE_SIMPLE': ('eur2', None, None, None),
    'COSTE_DOBLE': ('eur2', None, None, None),
    'FC_SIMPLE': ('pct', None, None, None),
    'FC_DOBLE': ('pct', None, None, None),
    'FC_PONDERADO': ('pct', None, None, None),
    'COSTE_RACION_PATATA_CONG': ('eur2', None, None, None),
    'COSTE_RACION_PATATA_FRESCA': ('eur2', None, None, None),
    'ACEITE_RACION': ('eur2', None, None, None),
    'MARGEN_SALA_MENU': ('eur2', None, None, None),
    'MARGEN_PLATAFORMA_MENU': ('eur2', None, None, None),
}
#: Canales del caso que no están en el libro (F1-camino §2, sobre R §3): comensales de sala al día y pedidos al día para
#: recoger y de plataforma (1 pedido = 1,7 comensales equivalentes), para la FAQ y el Word.
COMENSALES_SALA = 52.7
PEDIDOS_LLEVAR_DIA = 12.9
PEDIDOS_PLATAFORMA_DIA = 12.0
EQUIVALENTES_PEDIDO = 1.7


def _fmt_calc(tok):
    """Formato de un token de la calculadora que no está en `CALC` (por su nombre)."""
    if tok.startswith(('FC_', 'RENDIMIENTO', 'PCT_', 'MARGEN_PCT')):
        return 'pct'
    if tok.startswith(('COSTE', 'PRECIO', 'MARGEN', 'PVP')):
        return 'eur2'
    return 'int'


def fmt(v, f):
    return _fmt(v, f).replace('-', '−', 1) if v < 0 else _fmt(v, f)   # signo menos tipográfico


def _r(v, nd=0):
    """Redondeo comercial (mitad hacia arriba, como Excel) y no el bancario de `round()`: 4,975 → 4,98 y 27.940,5 →
    27.941 (revisión final de la Pizzería, m2). `round(v, 9)` quita el ruido binario del float antes de pasar a Decimal."""
    q = Decimal(1).scaleb(-nd)
    return Decimal(str(round(v, 9))).quantize(q, rounding=ROUND_HALF_UP)


def _fmt(v, f):
    if f in ('int', 'int_up'):
        n = math.ceil(v - 1e-9) if f == 'int_up' else _r(v)
        return '{:,}'.format(int(n)).replace(',', '.')
    if f == 'eur':
        return '{:,}'.format(int(_r(v))).replace(',', '.') + ' €'
    if f == 'eur2':
        return '{:,.2f}'.format(_r(v, 2)).replace(',', 'X').replace('.', ',').replace('X', '.') + ' €'
    if f in ('pct', 'pct0'):
        txt = '{:f}'.format(_r(v * 100, 1 if f == 'pct' else 0)).replace('.', ',')
        return txt + ' %'
    if f in ('dec1', 'dec2'):
        return '{:f}'.format(_r(v, 1 if f == 'dec1' else 2)).replace('.', ',')
    raise ValueError(f)


def _fila(ws, rx, col=1):
    for r in range(1, ws.max_row + 1):
        a = ws.cell(r, col).value
        if isinstance(a, str) and re.search(rx, a.strip(), re.I):
            return r
    return None


def leer(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    crudo, falta, textos = {}, [], {}
    for tok, (hoja, rx, col, f) in MAPA.items():
        ws = wb[hoja]
        r = _fila(ws, rx)
        v = ws.cell(r, col).value if r else None
        if isinstance(v, str) and tok.startswith('PAYBACK') and v.strip():   # «Más de 3 años»: texto del libro
            textos[tok] = v.strip()
            continue
        if not isinstance(v, (int, float)):
            falta.append('%s (%s · %s): %r' % (tok, hoja, rx, v))
            continue
        crudo[tok] = (float(v), f)
    # derivados
    if all(k in crudo for k in ('CAPEX', 'FIANZA', 'RENTA_PREVIA')):    # las 18 partidas + imprevistos (R §3: 104.350)
        crudo['INVERSION_PARTIDAS'] = (crudo['CAPEX'][0] - crudo['FIANZA'][0] - crudo['RENTA_PREVIA'][0], 'eur')
    if all(k in crudo for k in ('COMENSALES_DIA', 'BE_COMENSALES_DIA')):    # holgura sobre el equilibrio CONTABLE
        ws = wb[PE]
        be = float(ws.cell(_fila(ws, MAPA['BE_COMENSALES_DIA'][1]), 2).value)
        crudo['HOLGURA_CONTABLE'] = (crudo['COMENSALES_DIA'][0] / be - 1, 'pct')
    if 'MES_SALDO_MINIMO' in crudo:
        textos['MES_SALDO_MINIMO_NOMBRE'] = MESES[int(round(crudo['MES_SALDO_MINIMO'][0])) - 1]
    if all(k in crudo for k in ('ORIGEN_FONDOS', 'NECESIDAD_CAJA')):
        crudo['EXCESO_FINANCIACION'] = (crudo['ORIGEN_FONDOS'][0] - crudo['NECESIDAD_CAJA'][0], 'eur')
    for i in (1, 2, 3):                                                   # el P&L no tiene fila de EBITDA
        ks = ('RAI_A%d' % i, 'AMORTIZACION_A%d' % i, 'INTERESES_A%d' % i)
        if all(k in crudo for k in ks):
            crudo['EBITDA_A%d' % i] = (sum(crudo[k][0] for k in ks), 'eur')
            if 'VENTAS_A%d' % i in crudo:
                crudo['EBITDA_PCT_A%d' % i] = (crudo['EBITDA_A%d' % i][0] / crudo['VENTAS_A%d' % i][0], 'pct')
    if all(k in crudo for k in ('VENTAS_A1', 'DIAS_APERTURA')):
        crudo['VENTAS_DIA_A1'] = (crudo['VENTAS_A1'][0] / crudo['DIAS_APERTURA'][0], 'eur')
    if all(k in crudo for k in ('VENTAS_A1', 'DELIVERY_PCT')):
        crudo['VENTAS_PLATAFORMA_A1'] = (crudo['VENTAS_A1'][0] * crudo['DELIVERY_PCT'][0], 'eur')
    if all(k in crudo for k in ('BEBIDA_21_PCT', 'PCT_BEBIDA')):          # parte de las ventas al 21 % (R §3: 4,5 %)
        crudo['VENTAS_21_PCT'] = (crudo['PCT_BEBIDA'][0] * crudo['BEBIDA_21_PCT'][0], 'pct')
    crudo['COMENSALES_SALA'] = (float(COMENSALES_SALA), 'int')            # F1-camino §2 (no está en el libro)
    crudo['PEDIDOS_LLEVAR_DIA'] = (PEDIDOS_LLEVAR_DIA, 'int')
    crudo['PEDIDOS_PLATAFORMA_DIA'] = (PEDIDOS_PLATAFORMA_DIA, 'int')
    if 'COMENSALES_DIA' in crudo:                                          # ≈ 25 pedidos al día (F1-camino §2)
        crudo['PEDIDOS_DIA'] = ((crudo['COMENSALES_DIA'][0] - COMENSALES_SALA) / EQUIVALENTES_PEDIDO, 'int')
    if 'AFORO' in crudo:
        crudo['ROT_SALA'] = (COMENSALES_SALA / crudo['AFORO'][0], 'dec1')  # la rotación real de la sala (≈ 1,5)
    if all(k in crudo for k in ('VENTAS_A1', 'PCT_COMIDA')):               # sensibilidad del Word: 1 punto de comida
        crudo['PUNTO_FOOD_COST_A1'] = (crudo['VENTAS_A1'][0] * crudo['PCT_COMIDA'][0] * 0.01, 'eur')
    return crudo, falta, textos


def leer_calc(path):
    """Tokens de la calculadora (opcional, no entra en los gates). Si el fichero no existe, avisa y sigue."""
    if not os.path.exists(path):
        return {}, ['la calculadora %s aún no existe: se omiten sus tokens' % path]
    try:
        expuestos = dict(importlib.import_module('calculadora_ham_es').EXPUESTOS)
    except Exception as e:                                                  # noqa: BLE001
        print('  AVISO calculadora_ham_es no se importa (%s): sin tokens de la calculadora' % e.__class__.__name__)
        expuestos = {}
    wb = openpyxl.load_workbook(path, data_only=True)
    crudo, falta = {}, []
    todos = dict(CALC)
    for tok in expuestos:
        todos.setdefault(tok, (_fmt_calc(tok), None, None, None))
    for tok, (f, hoja, rx_fila, rx_col) in todos.items():
        v = None
        if tok in expuestos:
            h, c = expuestos[tok]
            v = wb[h][c].value if h in wb.sheetnames else None
        elif hoja and hoja in wb.sheetnames:
            ws = wb[hoja]
            r = _fila(ws, rx_fila)
            col = 2
            if rx_col:
                col = next((c.column for row in ws.iter_rows() for c in row
                            if isinstance(c.value, str) and re.search(rx_col, c.value.strip(), re.I)), None)
            v = ws.cell(r, col).value if (r and col) else None
        if isinstance(v, (int, float)):
            crudo[tok] = (float(v), f)
        else:
            falta.append('%s (calculadora · %s): %r' % (tok, expuestos.get(tok, hoja), v))
    return crudo, falta


def prueba_iva_bebida(path):
    """`PyG!G` de la línea de bebida = H × 21 % + (1 − H) × IVA de restauración (10 %), en valor y en fórmula."""
    wb_v = openpyxl.load_workbook(path, data_only=True)
    wb_f = openpyxl.load_workbook(path)
    ws_v, ws_f = wb_v[PYG], wb_f[PYG]
    r = _fila(ws_v, MAPA['VENTAS_BEBIDA_A1'][1])
    g, h = ws_v.cell(r, 7).value, ws_v.cell(r, 8).value
    form = str(ws_f.cell(r, 7).value)
    red = wb_v[SUP].cell(_fila(wb_v[SUP], r'^iva repercutido de restauraci'), 2).value
    ok = (isinstance(g, (int, float)) and isinstance(h, (int, float)) and isinstance(red, (int, float))
          and abs(g - (h * 0.21 + (1 - h) * red)) < 1e-9 and abs(red - 0.10) < 1e-9 and '*0.21' in form
          and ('H%d' % r) in form)
    return ok, 'PyG!G%d = %s (H%d = %s, reducido %s) · fórmula %s' % (r, g, r, h, red, form[:90])


def prueba_mes1_enero(path, v):
    """F1-camino §7 F2 (2): pagas 7/12 en `Supuestos!B60/B61`, la tupla de enero en `Tesorería!B5:M5` y el mes del
    saldo mínimo coherente con la fila del saldo acumulado."""
    import contenido_plan_negocio_hamburgueseria.a as A                      # noqa: E402
    wb = openpyxl.load_workbook(path, data_only=True)
    sup, tes = wb[SUP], wb[TES]
    pagas = (sup['B60'].value, sup['B61'].value)
    ok_pagas = pagas == (7, 12) and re.search(r'paga extra de verano', str(sup['A60'].value), re.I) is not None
    r_est = _fila(tes, r'^estacionalidad del mes')
    fila = [tes.cell(5, c).value for c in range(2, 14)]
    ok_tupla = (r_est == 5 and all(isinstance(x, (int, float)) for x in fila)
                and all(abs(x - y) < 1e-6 for x, y in zip(fila, A.ESTACIONALIDAD)))
    r_saldo = _fila(tes, MAPA['SALDO_MES12'][1])
    saldos = [tes.cell(r_saldo, c).value for c in range(2, 14)] if r_saldo else []
    mes = int(round(v.get('MES_SALDO_MINIMO', 0)))
    ok_mes = (len(saldos) == 12 and all(isinstance(x, (int, float)) for x in saldos)
              and 1 <= mes <= 12 and saldos.index(min(saldos)) + 1 == mes
              and abs(min(saldos) - v.get('SALDO_MINIMO', 0)) < 0.01)
    return [
        ('Supuestos!B60/B61 = 7/12 (pagas con el Mes 1 en enero): %s' % (pagas,), ok_pagas),
        ('Tesorería!B5:M5 = la tupla de enero de a.ESTACIONALIDAD (fila %s)' % r_est, ok_tupla),
        ('mes del saldo mínimo = el de la fila del saldo acumulado: Mes %d (%s)' % (mes, MESES[mes - 1] if 1 <= mes <= 12
                                                                                   else '?'), ok_mes),
    ]


#: Restos de la semilla (Sushi Bar) y de sus moldes (Taquería, Pollería, Pizzería) en el plan (F1-camino §7 F2 (4)).
RX_RESTOS = (r'sushi|sashimi|nigiri|\bmakis?\b|itamae|sushiman|\bshari\b|\bnori\b|wasabi|\bsake\b|gyoza|tempura|'
             r'anisakis|hangiri|arrocera|ph-metro|japon[eé]s|salm[oó]n|\bat[uú]n\b|combinados?|taquer|trompo|\btacos?\b|'
             r'poller|pollo asado|pizz|nixtamal|mezcal|tequila|aguas? frescas?|al pastor|guacamole')


def prueba_hamburgueseria(path, v, textos_payback):
    """F1-camino §3 y §7 F2 (3)-(4): rotación encendida (techo 3,5, caso 2,5-2,8), aforo 36, plataforma < 25 %, 0
    restos de la semilla y de los moldes, rótulos «Comensales/día» y payback como número."""
    wb = openpyxl.load_workbook(path)
    textos = [(ws.title + '!' + c.coordinate, c.value) for ws in wb.worksheets for row in ws.iter_rows() for c in row
              if isinstance(c.value, str)]
    restos = [k for k, t in textos if re.search(RX_RESTOS, t, re.I)]
    com = [k for k, t in textos if re.match(r'comensales/d[ií]a', t.strip(), re.I)]
    return [
        ('rotación: techo 3,5 y caso 2,5-2,8 (%s · %s)' % (v.get('ROT_MAX'), v.get('ROTACION_IMPLICITA')),
         abs(v.get('ROT_MAX', 0) - 3.5) < 1e-9 and 2.5 <= v.get('ROTACION_IMPLICITA', 0) <= 2.8),
        ('aforo = 36: %s' % v.get('AFORO'), abs(v.get('AFORO', 0) - 36) < 1e-9),
        ('plataforma < 25 %% de las ventas: %s' % v.get('DELIVERY_PCT'), v.get('DELIVERY_PCT', 1) < 0.25),
        ('0 restos del sushi bar, la taquería, la pollería o la pizzería en el plan: %s' % (restos[:6] or 'ninguno'),
         not restos),
        ('rótulos «Comensales/día» presentes: %d' % len(com), len(com) >= 2),
        ('payback antes de la deuda como número: %s' % (v.get('PAYBACK') if 'PAYBACK' in v
                                                       else textos_payback.get('PAYBACK')),
         'PAYBACK' in v and 'PAYBACK' not in textos_payback),
    ]


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--xlsx', required=True)
    ap.add_argument('--calc')
    ap.add_argument('--salida', default=SALIDA)
    a = ap.parse_args()
    crudo, falta, textos = leer(a.xlsx)
    falta_calc = []
    if a.calc:
        c2, falta_calc = leer_calc(a.calc)
        crudo.update(c2)
    tokens = dict((k, fmt(v, f)) for k, (v, f) in crudo.items())
    tokens.update(textos)
    json.dump({'fuente': os.path.basename(a.xlsx),
               'calculadora': os.path.basename(a.calc) if (a.calc and os.path.exists(a.calc)) else None,
               'tokens': tokens, 'valores': dict((k, v) for k, (v, _f) in crudo.items())},
              open(a.salida, 'w', encoding='utf-8'), ensure_ascii=False, indent=1, sort_keys=True)
    v = dict((k, x[0]) for k, x in crudo.items())
    iva_ok, iva_txt = prueba_iva_bebida(a.xlsx)
    rot_be = [v.get(k, 99) for k in ('ROT_BE_CAJA', 'ROT_BE_CAJA_A2', 'ROT_BE_CAJA_A3')]
    gates = [
        ('DSCR mínimo ≥ 1,25', v.get('DSCR_MIN', 0) >= 1.25),
        ('saldo mínimo de tesorería > 0', v.get('SALDO_MINIMO', -1) > 0),
        ('cobertura de horas ≥ 98 %', v.get('COBERTURA_HORAS', 0) >= 0.98),
        ('holgura sobre el equilibrio de caja del año 1 ≥ 15 %', v.get('HOLGURA_CAJA', 0) >= 0.15),
        ('holgura de caja del año 2 (sin carencia) ≥ 0 % (informativa)', v.get('HOLGURA_CAJA_A2', -1) >= 0),
        ('holgura de caja del año 3 ≥ 0 % (informativa)', v.get('HOLGURA_CAJA_A3', -1) >= 0),
        ('exceso de financiación (origen - usos) ≥ 0 y < 5 %', 0 <= v.get('DIFERENCIA_FONDOS_PCT', -1) < 0.05),
        ('resultado neto a régimen (años 2 y 3) entre el 7 y el 12 %',
         all(0.07 <= v.get('NETO_PCT_A%d' % i, 1) <= 0.12 for i in (2, 3))),
        ('resultado neto del año 1 ≤ 12 %', v.get('NETO_PCT_A1', 1) <= 0.12),
        ('rotación de equilibrio de caja ≤ techo del local (años 1-3)', max(rot_be) <= v.get('ROT_MAX', 0)),
        ('ventas por plataforma < 25 % (demo §2.11.9)', v.get('DELIVERY_PCT', 1) < 0.25),
        ('IVA de la bebida = H × 21 %% + (1 − H) × 10 %% (%s)' % iva_txt, iva_ok),
    ] + prueba_mes1_enero(a.xlsx, v) + prueba_hamburgueseria(a.xlsx, v, textos)
    for tok in ('COMENSALES_DIA', 'TICKET', 'TICKET_IVA', 'DIAS_APERTURA', 'VENTAS_A1', 'VENTAS_A2', 'VENTAS_A3',
                'MARGEN_BRUTO_PCT_A1', 'MERCANCIA_PCT_A1', 'PERSONAL_A1', 'PERSONAL_PCT_A1', 'ALQUILER_PCT_A1',
                'FIJOS_A1', 'EBITDA_A1', 'EBITDA_PCT_A1', 'RAI_A1', 'NETO_A1', 'NETO_A2', 'NETO_A3', 'NETO_PCT_A1',
                'NETO_PCT_A2', 'NETO_PCT_A3', 'INVERSION_PARTIDAS', 'CAPEX', 'FONDO_MANIOBRA', 'INVERSION_TOTAL',
                'IVA_INVERSION', 'NECESIDAD_CAJA', 'RECURSOS_PROPIOS', 'PRESTAMO', 'DIFERENCIA_FONDOS',
                'DIFERENCIA_FONDOS_PCT', 'CUOTA_ANUAL', 'DSCR_A2', 'DSCR_A3', 'DSCR_MIN', 'SALDO_MINIMO',
                'MES_SALDO_MINIMO', 'MES_SALDO_MINIMO_NOMBRE', 'SALDO_MES12', 'BE_COMENSALES_DIA',
                'BE_COMENSALES_DIA_CAJA', 'BE_COMENSALES_DIA_CAJA_A2', 'BE_COMENSALES_DIA_CAJA_A3',
                'HOLGURA_CONTABLE', 'HOLGURA_CAJA', 'HOLGURA_CAJA_A2', 'HOLGURA_CAJA_A3', 'ROTACION_IMPLICITA',
                'ROT_BE_CAJA', 'ROT_MAX', 'COBERTURA_HORAS', 'JORNADAS', 'PAYBACK', 'PAYBACK_CAPEX',
                'IVA_BEBIDA_LINEA', 'BEBIDA_21_PCT', 'VENTAS_21_PCT', 'IVA_COMPRA_COMIDA', 'IVA_COMPRA_BEBIDA',
                'ESC_PES_VENTAS', 'ESC_PES_NETO', 'ESC_OPT_VENTAS', 'ESC_OPT_NETO', 'COMENSALES_SALA', 'ROT_SALA',
                'PEDIDOS_DIA', 'PEDIDOS_LLEVAR_DIA', 'PEDIDOS_PLATAFORMA_DIA', 'PUNTO_FOOD_COST_A1') + tuple(CALC):
        print('  %-28s %s' % (tok, tokens.get(tok, '—')))
    for nombre, ok in gates:
        print('  [%s] %s' % ('OK' if ok else 'ROJO', nombre))
    for f in falta:
        print('  FALTA', f)
    for f in falta_calc:
        print('  FALTA (calculadora, no entra en el gate)', f)
    ok_total = all(ok for _n, ok in gates) and not falta
    print('cifras_caso_es.json: %d tokens · gates del caso D9 %s' % (len(tokens), 'VERDE' if ok_total else 'ROJA'))
    sys.exit(0 if ok_total else 1)


if __name__ == '__main__':
    main()
