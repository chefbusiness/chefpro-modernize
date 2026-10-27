#!/usr/bin/env python3
"""
contrato_wed.py — Wedding Venue Business Plan Kit (EN), SPEC D171: `wedding-venue-rental-agreement-templates.docx` «Venue
Rental Agreement Templates: Couples, Corporate Clients & Exclusive Caterer». Clon de `../plan-finca-eventos/
contrato_fin_es.py`: abre el docx ES del contrato de la finca como MOLDE FÍSICO (estilos `Normal` y `List Bullet`,
márgenes, fuentes), vacía el cuerpo (conserva la sección), pone la página en US Letter y escribe
`textos_en/contrato_en_1.json` + `textos_en/contrato_en_2.json` (tanda 1: `[{estilo, texto}]` con `**negrita**`).

Gates sobre el fichero escrito (D172): Letter, autor y docProps EN, 0 tokens, negrita cerrada, cabecera legal literal,
0 español, 0 no latinos, 0 `cifras_prohibidas` / `RX_WED_NO`, ningún «non-refundable» sin la tabla, la Cooling-Off Rule
nunca afirmada (decisión de orquestador 2) y estilos dentro del molde. Con `--pdf` (LibreOffice) cuenta las páginas.
Corre en el VPS (python-docx). Sesión Claude Code, 27-oct-2026 (F2-EN, paso 3).

    python3 bp9.py contrato --salida <dir> [--pdf]
"""
import argparse
import datetime
import json
import os
import re
import subprocess
import sys
import tempfile

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor

AQUI = os.path.dirname(os.path.abspath(__file__))
PD = os.path.dirname(AQUI)
REPO = os.path.abspath(os.path.join(AQUI, '..', '..', '..'))
sys.path.insert(0, os.path.join(PD, 'business-plans-en'))
sys.path.insert(0, AQUI)
os.environ.setdefault('BP_CONJUNTO', 'wed')
import mapas as M                                                                  # noqa: E402
from check_contrato import CABECERA, RX_COOLING_SI, RX_NONREF, RX_TOKEN            # noqa: E402

SLUG = 'wedding-venue-business-plan'
PRODUCTO = 'Wedding Venue Business Plan Kit'
URL = 'aichef.pro/en/digital-products/' + SLUG
FICHERO = 'wedding-venue-rental-agreement-templates.docx'
TITULO = 'Venue Rental Agreement Templates: Couples, Corporate Clients & Exclusive Caterer'
MOLDE = os.path.join(REPO, 'astro-site', 'public', 'dl', 'plan-negocio-finca-eventos-bodas',
                     'modelo-contrato-reserva-finca-eventos.docx')
TEXTOS = [os.path.join(AQUI, 'textos_en', 'contrato_en_%d.json' % n) for n in (1, 2)]
FECHA = datetime.datetime(2026, 10, 27)
#: formato directo del molde ES: (fuente, negrita, cursiva, color, puntos)
FORMATO = {'Kicker': (None, True, False, 'B08D57', 9), 'Title': ('Playfair Display', True, False, '1F1B16', 28),
           'Subtitle': (None, False, True, '6B5E4D', 12), 'H1': ('Playfair Display', True, False, '1F1B16', 18),
           'H2': ('Inter', True, False, '1F1B16', 13), 'Separador': (None, False, False, 'B08D57', 11)}
ESTILOS = set(FORMATO) | {'Normal', 'List Bullet', 'Firma'}
RX_NEGRITA = re.compile(r'(\*\*.+?\*\*)')


def runs(p, texto, fmt=None):
    for trozo in RX_NEGRITA.split(texto):
        if not trozo:
            continue
        negrita = trozo.startswith('**') and trozo.endswith('**') and len(trozo) > 4
        r = p.add_run(trozo[2:-2] if negrita else trozo)
        if fmt:
            fuente, b, i, color, pt = fmt
            if fuente:
                r.font.name = fuente
            r.font.bold, r.font.italic, r.font.size = b or negrita or None, i or None, Pt(pt)
            r.font.color.rgb = RGBColor.from_string(color)
        elif negrita:
            r.font.bold = True


def generar(salida, pdf=False):
    items = []
    for p in TEXTOS:
        items += json.load(open(p, encoding='utf-8'))
    err = []
    doc = Document(MOLDE)
    body = doc.element.body
    for hijo in list(body):
        if not hijo.tag.endswith('}sectPr'):
            body.remove(hijo)
    for it in items:
        est, t = it['estilo'], it['texto']
        if est not in ESTILOS:
            err.append('estilo desconocido %r' % est)
            continue
        if est in ('Normal', 'List Bullet', 'Firma'):
            p = doc.add_paragraph(style='List Bullet' if est == 'List Bullet' else 'Normal')
            p.paragraph_format.space_after = Pt(4 if est == 'List Bullet' else 8)
            runs(p, t)
        else:
            p = doc.add_paragraph(style='Normal')
            if est == 'Separador':
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(10 if est in ('H1', 'H2') else 0)
            runs(p, t or '·  ·  ·', FORMATO[est])
            if est in ('H1', 'H2', 'Kicker'):
                p.paragraph_format.keep_with_next = True
    for s in doc.sections:
        s.page_width, s.page_height = Inches(8.5), Inches(11)                     # US Letter (márgenes del molde)
    pr = doc.core_properties
    pr.title, pr.subject = '%s · %s' % (TITULO, PRODUCTO), '%s · v%s' % (PRODUCTO, M.VERSION)
    pr.author, pr.last_modified_by = 'AI Chef Pro', 'AI Chef Pro'
    pr.comments, pr.keywords = URL, 'venue rental agreement, wedding venue contract, event venue contract template, AI Chef Pro'
    pr.category, pr.language = 'AI Chef Pro · Digital products', 'en-US'
    pr.created = pr.modified = FECHA
    pr.revision = 1
    os.makedirs(salida, exist_ok=True)
    destino = os.path.join(salida, FICHERO)
    doc.save(destino)
    d2 = Document(destino)
    pars = [p.text for p in d2.paragraphs]
    todo = '\n'.join(pars)
    s = d2.sections[0]
    if abs(s.page_width - Inches(8.5)) > Inches(0.01) or abs(s.page_height - Inches(11)) > Inches(0.01):
        err.append('no es US Letter: %.2f x %.2f in' % (s.page_width.inches, s.page_height.inches))
    if CABECERA.replace('**', '') not in pars[:8]:
        err.append('falta la cabecera legal literal en los 8 primeros párrafos')
    if RX_TOKEN.search(todo):
        err.append('tokens sin resolver: %r' % RX_TOKEN.findall(todo)[:5])
    if '**' in todo:
        err.append('negrita sin cerrar (**)')
    for i, t in enumerate(pars):
        r = [x for x in M.restos_espanol(t) if not x[0].startswith('oficio')]
        if r:
            err.append('[%d] español: %s «%s»' % (i, r[0][0], t[:60]))
        if M.no_latinos(t):
            err.append('[%d] caracteres no latinos' % i)
        for motivo, frag in M.cifras_prohibidas(t):
            err.append('[%d] %s «%s»' % (i, motivo, frag))
        if RX_NONREF.search(t) and not re.search(r'(?i)schedule|table|clause|section', t):
            err.append('[%d] «non-refundable» sin la tabla de cancelación' % i)
        for m in RX_COOLING_SI.finditer(t):
            if 'whether' not in t[max(0, m.start() - 60):m.start()].lower():
                err.append('[%d] afirma que la FTC Cooling-Off Rule aplica' % i)
    cp = d2.core_properties
    if cp.author != 'AI Chef Pro' or cp.created.year != 2026 or cp.comments != URL or not cp.title.startswith(TITULO):
        err.append('docProps: autor, año, URL o título')
    estilos = {p.style.name for p in d2.paragraphs}
    if not estilos <= {'Normal', 'List Bullet'}:
        err.append('estilos fuera del molde: %r' % estilos)
    paginas = None
    if pdf:
        tmp = tempfile.mkdtemp(prefix='wed-contrato-')
        subprocess.run(['soffice', '--headless', '--convert-to', 'pdf', '--outdir', tmp, destino],
                       stdout=subprocess.PIPE, stderr=subprocess.PIPE, check=False)
        pp = os.path.join(tmp, FICHERO.replace('.docx', '.pdf'))
        if os.path.exists(pp):
            info = subprocess.run(['pdfinfo', pp], stdout=subprocess.PIPE, universal_newlines=True).stdout
            m = re.search(r'Pages:\s+(\d+)', info)
            paginas = int(m.group(1)) if m else None
            if 'Page size:      612 x 792' not in info:
                err.append('el PDF no sale en Letter (%s)' % re.search(r'Page size:.*', info).group(0))
        else:
            err.append('LibreOffice no generó el PDF')
    for e in err:
        print('ERROR', e)
    print('%s: %d párrafos · %d palabras · Letter · estilos %s%s · %s' % (
        FICHERO, len(d2.paragraphs), len(todo.split()), sorted(estilos),
        '' if paginas is None else ' · %d páginas (PDF)' % paginas, 'VERDE' if not err else 'ROJO (%d errores)' % len(err)))
    return destino, 1 if err else 0


if __name__ == '__main__':
    ap = argparse.ArgumentParser()
    ap.add_argument('--salida', required=True)
    ap.add_argument('--pdf', action='store_true', help='convierte con LibreOffice y cuenta las páginas')
    a = ap.parse_args()
    sys.exit(generar(a.salida, a.pdf)[1])
