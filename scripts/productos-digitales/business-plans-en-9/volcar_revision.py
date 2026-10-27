#!/usr/bin/env python3
"""
volcar_revision.py — Wedding Venue Business Plan Kit (EN) · vuelca a TEXTO los 5 entregables de
`astro-site/public/dl/wedding-venue-business-plan/` en `revision/<fichero>.txt` para la revisión D174: xlsx hoja a hoja
y celda a celda (valor en caché; las fórmulas, marcadas con «=» y su resultado), docx párrafo a párrafo con su estilo.
Corre en el VPS (openpyxl + python-docx). Sesión Claude Code, 27-oct-2026 (F2-EN, paso 4).

    python3 volcar_revision.py
"""
import datetime
import os

import docx
import openpyxl

AQUI = os.path.dirname(os.path.abspath(__file__))
REPO = os.path.abspath(os.path.join(AQUI, '..', '..', '..'))
DIR = os.path.join(REPO, 'astro-site', 'public', 'dl', 'wedding-venue-business-plan')
SALIDA = os.path.join(AQUI, 'revision')


def fmt(v):
    if isinstance(v, float):
        return ('%.4f' % v).rstrip('0').rstrip('.')
    if isinstance(v, (datetime.date, datetime.datetime)):
        return v.strftime('%m/%d/%Y')
    return str(v)


def xlsx(p):
    wf, wv = openpyxl.load_workbook(p), openpyxl.load_workbook(p, data_only=True)
    out = ['# %s' % os.path.basename(p), 'docProps: title=%r · subject=%r · creator=%r' % (
        wf.properties.title, wf.properties.subject, wf.properties.creator)]
    for ws in wf.worksheets:
        out += ['', '## %s' % ws.title]
        for row in ws.iter_rows():
            for c in row:
                if c.value is None:
                    continue
                v = wv[ws.title][c.coordinate].value
                if c.data_type == 'f':
                    out.append('%s = %s' % (c.coordinate, fmt(v)))
                else:
                    out.append('%s: %s' % (c.coordinate, fmt(v)))
        for dv in ws.data_validations.dataValidation:
            if dv.type == 'list' or dv.error:
                out.append('DV %s: %s %s · %s' % (dv.sqref, dv.type, dv.formula1 or '', dv.error or ''))
    return out


def docx_(p):
    d = docx.Document(p)
    cp = d.core_properties
    out = ['# %s' % os.path.basename(p), 'core: title=%r · author=%r · página %.1f × %.1f in' % (
        cp.title, cp.author, d.sections[0].page_width.inches, d.sections[0].page_height.inches)]
    for i, x in enumerate(d.paragraphs):
        if x.text.strip():
            out.append('[%d] (%s) %s' % (i, x.style.name, x.text))
    return out


def main():
    os.makedirs(SALIDA, exist_ok=True)
    for f in sorted(os.listdir(DIR)):
        p = os.path.join(DIR, f)
        lineas = xlsx(p) if f.endswith('.xlsx') else docx_(p) if f.endswith('.docx') else None
        if lineas is None:
            continue
        with open(os.path.join(SALIDA, f + '.txt'), 'w', encoding='utf-8') as fh:
            fh.write('\n'.join(lineas) + '\n')
        print('revision/%s.txt: %d líneas · %d palabras' % (f, len(lineas), sum(len(x.split()) for x in lineas)))


if __name__ == '__main__':
    main()
