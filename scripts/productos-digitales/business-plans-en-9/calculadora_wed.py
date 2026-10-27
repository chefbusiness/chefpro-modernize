#!/usr/bin/env python3
"""
calculadora_wed.py — Wedding Venue Business Plan Kit (EN), SPEC D168: `wedding-venue-pricing-calculator.xlsx` «Wedding
Venue Pricing & Quote Calculator (Season, Day, Retainer & Cancellation Schedule)». Es `../plan-finca-eventos/
calculadora_fin_es.py` traducida y adaptada a EE. UU.: mismas hojas (`Calculator` · `Examples` · `Instructions`), mismas
filas, helpers, `inject_cache` y autotest (pycel + caché = `calcular()`, 3 ejemplos + 3 casos propios). Cambios D168:
  1. canal Couple / Corporate (al corporativo se le presupuesta «plus sales tax»; a la pareja, con el impuesto incluido);
  2. tipologías Wedding $8,000 · Corporate event $3,500 · Private celebration $4,000 · Micro wedding / elopement $2,000
     (tarifa de un sábado de shoulder season), F&B del caterer por invitado $110 · $50 · $55 · $80 y coste directo $1,000 ·
     $500 · $500 · $300 (research §5.1 y §8: Kit estimate);
  3. temporadas Peak +15 % / Shoulder 0 / Off-peak −20 % con mínimo de invitados 100 / 75 / 50; meses US (may, jun, sep,
     oct = peak; abr, jul, ago, nov = shoulder; dic-mar = off-peak) y días Saturday 0 · Friday or Sunday −15 % · Weekday
     −30 %;
  4. extras: ceremony on site $750 · additional hour $500 · premium décor & lighting $800 · rehearsal the day before $400
     (sustituye al alojamiento);
  5. personal $18/h + payroll taxes 10 %; redondeo $50; caterer commission 10 % (supuesto); sales tax on venue charges 8 %
     y sales tax on the caterer's F&B 8 % (solo para el presupuesto orientativo de la pareja); sin «commission with tax»
     (en el caso US la comisión va al 0 %);
  6. cobros Couple 25 % / 35 % a 180 días / 40 % a 30 días · Corporate 30 / 30 a 60 / 40 a 15; tramos de cancelación = ES
     (10/25/50/75/100 % · 5/15/35/60/100 % a 12/6/3/1/0 meses) como «a reasonable estimate of the loss, not a penalty»
     con rebooking credit; sin citas del derecho español (decisiones de orquestador 2 y 4: nada de 16 CFR 429 ni §1671);
  7. formatos $#,##0.00 y mm/dd/yyyy, US Letter, docProps D20, versión 2.2 · October 2026 (como el libro y el Word).
Corre en el VPS (openpyxl + pycel). Sesión Claude Code, 27-oct-2026 (F2-EN, paso 1b).

    python3 bp9.py calculadora --salida <dir>               # genera + inject_cache + autotest
    python3 bp9.py calculadora --salida <dir> --autotest    # solo el autotest sobre el fichero ya generado
"""
import argparse
import datetime
import math
import os
import re
import sys
from collections import OrderedDict

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill, Protection
from openpyxl.worksheet.datavalidation import DataValidation

AQUI = os.path.dirname(os.path.abspath(__file__))
PD = os.path.dirname(AQUI)
sys.path.insert(0, os.path.join(PD, 'planes-v2_0'))
sys.path.insert(0, os.path.join(PD, 'business-plans-en'))
sys.path.insert(0, PD)
os.environ.setdefault('BP_CONJUNTO', 'wed')
import motor                                                                       # noqa: E402  (solo gate_formatos)
import mapas as M                                                                  # noqa: E402  (detectores y textos fijos)

SLUG = 'wedding-venue-business-plan'
PRODUCTO = 'Wedding Venue Business Plan Kit'
URL = 'aichef.pro/en/digital-products/' + SLUG
FICHERO = 'wedding-venue-pricing-calculator.xlsx'
TITULO = 'Wedding Venue Pricing & Quote Calculator (Season, Day, Retainer & Cancellation Schedule)'
KEYWORDS = ('wedding venue pricing, wedding venue cost calculator, wedding venue business plan, venue rental fee, '
            'AI Chef Pro')
VERSION, MES = M.VERSION, M.MES_VERSION
VERSION_LINEA = 'Version %s · %s · %s · info@aichef.pro' % (VERSION, MES, URL)
FOOTER = M.FOOTER_EN
H1, H2, H3 = 'Calculator', 'Examples', 'Instructions'
VERDE = PatternFill('solid', fgColor='E8F5E9')
GRIS = PatternFill('solid', fgColor='F2F2F2')
FUENTE = 'Calibri'
CANALES = ['Couple', 'Corporate']
USD, ENT, PCT, DEC, FECHA = '$#,##0.00', '#,##0', '0%', '#,##0.0', 'mm/dd/yyyy'
EPOCA = datetime.date(1899, 12, 30)


def serial(d):
    """Fecha → número de serie de Excel (así se guardan las fechas: pycel y Excel leen lo mismo)."""
    return (d - EPOCA).days


# ---------------------------------------------------------------------------------------------------------------------
# Tablas de datos (SPEC D168; research §5.1 y §8). FUENTE ÚNICA: la calculadora las escribe en celdas verdes.
# Todas las cifras son Kit estimates: el usuario pone las de su venue.
# ---------------------------------------------------------------------------------------------------------------------
#: tipología → (tarifa del venue un sábado de shoulder season, sin sales tax · F&B del caterer por invitado · coste directo)
TIPOS = OrderedDict([
    ('Wedding', (8000, 110.0, 1000)), ('Corporate event', (3500, 50.0, 500)),
    ('Private celebration', (4000, 55.0, 500)), ('Micro wedding / elopement', (2000, 80.0, 300))])
#: temporada → (ajuste sobre la tarifa, mínimo de invitados para el presupuesto)
TEMPORADAS = OrderedDict([('Peak', (0.15, 100)), ('Shoulder', (0.0, 75)), ('Off-peak', (-0.20, 50))])
#: tipo de día → ajuste sobre la tarifa
DIAS = OrderedDict([('Saturday', 0.0), ('Friday or Sunday', -0.15), ('Weekday', -0.30)])
#: mes → temporada (research §5.1-5.2: octubre, septiembre y junio son los picos; dic-mar, off-peak)
MESES = [('January', 'Off-peak'), ('February', 'Off-peak'), ('March', 'Off-peak'), ('April', 'Shoulder'),
         ('May', 'Peak'), ('June', 'Peak'), ('July', 'Shoulder'), ('August', 'Shoulder'), ('September', 'Peak'),
         ('October', 'Peak'), ('November', 'Shoulder'), ('December', 'Off-peak')]
#: día de la semana (Sunday = 1, como WEEKDAY de Excel) → tipo de día
SEMANA = [('Sunday', 'Friday or Sunday'), ('Monday', 'Weekday'), ('Tuesday', 'Weekday'), ('Wednesday', 'Weekday'),
          ('Thursday', 'Weekday'), ('Friday', 'Friday or Sunday'), ('Saturday', 'Saturday')]
#: servicio extra del venue → precio de EJEMPLO sin sales tax
EXTRAS = OrderedDict([
    ('cer', ('Ceremony on site (per ceremony)', 750.0)),
    ('hex', ('Additional hour of venue use (per hour)', 500.0)),
    ('mon', ('Premium decor & lighting package (per event)', 800.0)),
    ('reh', ('Rehearsal the day before (per rehearsal)', 400.0))])
#: política por canal: Couple · Corporate
CANAL_TAB = OrderedDict([
    ('pc_1', ('Payment 1: retainer to book the date (% of the total)', (0.25, 0.30),
              'At signing. The retainer is applied to the price; what you keep if the client cancels is set only by the '
              'cancellation schedule below.')),
    ('pc_2', ('Payment 2: second installment (% of the total)', (0.35, 0.30), 'Due on the date in the next row.')),
    ('pc_3', ('Payment 3: final balance before the event (% of the total)', (0.40, 0.40),
              'The rest, before the event. Issue a receipt for every payment, the retainer included.')),
    ('d_2', ('Second installment: days before the event', (180, 60),
             'Couples, about six months before; corporate clients, two months. If they book later, both are due together.')),
    ('d_3', ('Final balance: days before the event', (30, 15), 'Couples, 30 days before; corporate clients, 15.')),
    ('cn1', ('Cancellation tier 1: what you keep (% of the total)', (0.10, 0.05),
             'Furthest from the event. Never more than the retainer you collected.')),
    ('cn2', ('Cancellation tier 2: what you keep (% of the total)', (0.25, 0.15), 'Rises as the event gets closer.')),
    ('cn3', ('Cancellation tier 3: what you keep (% of the total)', (0.50, 0.35),
             'Each tier leaves less time to rebook the date.')),
    ('cn4', ('Cancellation tier 4: what you keep (% of the total)', (0.75, 0.60), 'The last months.')),
    ('cn5', ('Cancellation tier 5: what you keep (% of the total)', (1.0, 1.0),
             'Last month or no-show: 100%. The agreement credits back what you recover by rebooking the date and the '
             'costs you save by not holding the event.')),
])
#: tramos de cancelación: meses de antelación a partir de los cuales aplica cada uno (de más lejano a más cercano, sin huecos)
TRAMOS = [12, 6, 3, 1, 0]

ENTRADAS = [
    ('canal', 'Client', '@', 'Couple',
     'Couple or Corporate: they change the payment schedule and the cancellation tiers (table at the end). A corporate '
     'client gets the quote «plus sales tax»; a couple, with sales tax included.'),
    ('tipo', 'Event type', '@', 'Wedding',
     'Wedding, corporate event (meeting, dinner, holiday party), private celebration (shower, milestone birthday, rehearsal '
     'dinner) or micro wedding / elopement. It sets the venue rate, the caterer\'s food & beverage per guest and the direct '
     'cost from the tables at the end.'),
    ('fecha', 'Event date', FECHA, serial(datetime.date(2027, 10, 9)),
     'Type the full date (mm/dd/yyyy). It sets the season and the type of day from the month and weekday tables.'),
    ('guests', 'Expected guests', ENT, 130, 'Whole number of guests. The quote uses the higher of this number and the '
                                            'season minimum.'),
    ('n_cer', 'Ceremonies on site (no.)', ENT, 1, 'At the price in the extra services table.'),
    ('n_hex', 'Additional hours of venue use (no.)', ENT, 2, 'Hours beyond the base rental window in your agreement.'),
    ('n_mon', 'Premium decor & lighting packages (no.)', ENT, 1, 'Decor, furniture and lighting beyond the standard setup.'),
    ('n_reh', 'Rehearsals the day before (no.)', ENT, 0, '0 if the couple does not book a rehearsal at the venue.'),
    ('h_aux', 'Venue staff hours on the event (coordinator and attendants)', DEC, 16,
     'Kit estimate: type your team\'s real hours (servers, bartenders and kitchen staff belong to the caterer).'),
    ('wage', 'Hourly wage of venue staff', USD, 18.0, 'Kit estimate. At least your state or local minimum wage.'),
    ('ss', 'Employer payroll taxes (% of wages)', PCT, 0.10,
     'FICA 7.65% plus federal and state unemployment taxes, about 10% (Kit estimate): check with your payroll provider.'),
    ('cost_ext', 'Direct cost of extra services (% of their price)', PCT, 0.40,
     'Kit estimate: what it costs you to deliver the ceremony, the decor package, the extra hours and the rehearsal.'),
    ('redondeo', 'Rounding of the venue rental fee (multiple, in dollars)', ENT, 50,
     'The venue rental fee is rounded to the nearest multiple of this amount.'),
    ('canon', 'Caterer commission (% of the caterer\'s food & beverage sales)', PCT, 0.10,
     'An editable ASSUMPTION, not a market figure: negotiate it with your caterer. For reference only: the caterer pays it '
     'to the venue, not the client.'),
    ('iva', 'Sales tax on venue charges (example)', PCT, 0.08,
     'Example rate (state + local). Whether a venue rental fee is taxable depends on the state: ask your state\'s '
     'department of revenue.'),
    ('iva_menu', 'Sales tax on the caterer\'s food & beverage (reference only)', PCT, 0.08,
     'The caterer bills the menu and the bar with its own sales tax, service charge and gratuity; they are not venue '
     'revenue.'),
]
K = OrderedDict((k, 5 + i) for i, (k, *_r) in enumerate(ENTRADAS))
F_CALC = 5 + len(ENTRADAS) + 1

# nombre, rótulo, formato, expresión, nota. {x}: celda de la entrada/salida; {t_xxx}: búsquedas en las tablas
SALIDAS = [
    ('temporada', 'Season of the date (from the month table)', 'General', '{t_temporada}',
     'Looks up the month of the date in the table at the end.'),
    ('dia', 'Type of day (from the weekday table)', 'General', '{t_dia}',
     'Looks up the weekday of the date in the table at the end.'),
    ('aj_t', 'Season adjustment (premium + / discount −)', PCT, '{t_aj_t}', 'From the season table.'),
    ('aj_d', 'Day adjustment (premium + / discount −)', PCT, '{t_aj_d}', 'From the day table.'),
    ('tarifa', 'Venue rate for the event type (Saturday, shoulder season)', USD, '{t_tarifa}', 'From the event type table.'),
    ('espacio', 'VENUE RENTAL FEE EXCL. SALES TAX', USD, 'ROUND({tarifa}*(1+{aj_t})*(1+{aj_d})/{redondeo},0)*{redondeo}',
     'Rate × (1 + season adjustment) × (1 + day adjustment), rounded to the chosen multiple.'),
    ('min_inv', 'Season minimum (guests)', ENT, '{t_min}', 'From the season table.'),
    ('garant', 'Guaranteed guest count', ENT, 'MAX({guests},{min_inv})',
     'The higher of the expected guests and the season minimum.'),
    ('aviso', 'Minimum check', 'General',
     'IF({guests}<{min_inv},"Below the minimum: the quote uses the season minimum","Meets the minimum")', ''),
    ('e_cer', 'Ceremonies', USD, '{n_cer}*{x_cer}', 'No. × price from the extra services table.'),
    ('e_hex', 'Additional hours', USD, '{n_hex}*{x_hex}', ''),
    ('e_mon', 'Premium decor & lighting', USD, '{n_mon}*{x_mon}', ''),
    ('e_reh', 'Rehearsals', USD, '{n_reh}*{x_reh}', ''),
    ('extras', 'Extra services excl. sales tax', USD, '{e_cer}+{e_hex}+{e_mon}+{e_reh}', ''),
    ('base', 'Taxable amount (venue rental + services)', USD, '{espacio}+{extras}', ''),
    ('iva_amt', 'Sales tax on venue charges', USD, '{base}*{iva}', ''),
    ('total', 'VENUE TOTAL (incl. sales tax)', USD, '{base}+{iva_amt}',
     'It does not include food or the bar: the exclusive caterer bills them separately.'),
    ('quoted', 'Amount you quote (Couple: incl. sales tax · Corporate: plus sales tax)', USD,
     'IF({canal}="Couple",{total},{base})',
     'A corporate client gets the quote excl. sales tax («plus sales tax»; if it is tax-exempt, ask for its exemption '
     'certificate); a couple, with sales tax included.'),
    ('eur_inv', 'Venue rental & services per guaranteed guest (excl. sales tax)', USD, '{base}/{garant}',
     'To compare with other venues.'),
    ('cb1', 'Payment 1: retainer', USD, '{total}*{t_pc_1}', 'At signing.'),
    ('cb2', 'Payment 2: second installment', USD, '{total}*{t_pc_2}', ''),
    ('f_2', 'Second installment due by', FECHA, '{fecha}-{t_d_2}', 'Event date minus the days in the table.'),
    ('cb3', 'Payment 3: final balance', USD, '{total}*{t_pc_3}', ''),
    ('f_3', 'Final balance due by', FECHA, '{fecha}-{t_d_3}', 'Event date minus the days in the table.'),
    ('menu_tot', 'Caterer\'s food & beverage excl. sales tax (estimate, guaranteed guests)', USD, '{t_menu}*{garant}',
     'FOR REFERENCE. Food & beverage per guest from the event type table × guaranteed guests. It is not venue revenue.'),
    ('canon_est', 'ESTIMATED CATERER COMMISSION (for reference)', USD, '{menu_tot}*{canon}',
     'The exclusive caterer pays it to the venue, not the client. It is an assumption: negotiate it.'),
    ('menu_cli', 'Caterer\'s food & beverage incl. sales tax (reference, billed by the caterer)', USD,
     '{menu_tot}*(1+{iva_menu})', 'So the client sees the full cost; the caterer\'s service charge and gratuity come on '
                                  'top.'),
    ('presup', 'Full reference budget for the client (venue incl. sales tax + food & beverage)', USD,
     '{total}+{menu_cli}', 'What the client pays the venue plus what it pays the caterer.'),
    ('c_fijo', 'Direct event cost: cleaning, restrooms, generator fuel, parking and security', USD, '{t_cfijo}',
     'From the event type table (Kit estimate).'),
    ('c_aux', 'Venue staff cost (wages + payroll taxes)', USD, '{h_aux}*{wage}*(1+{ss})', ''),
    ('c_ext', 'Direct cost of the extra services', USD, '{extras}*{cost_ext}', ''),
    ('c_total', 'Total direct cost of the event', USD, '{c_fijo}+{c_aux}+{c_ext}', ''),
    ('ingresos', 'Venue revenue from the event (excl. sales tax, commission included)', USD, '{base}+{canon_est}',
     'Venue rental + services + estimated commission.'),
    ('profit', 'Event contribution for the venue ($)', USD, '{ingresos}-{c_total}',
     'Before rent, salaried staff, utilities, insurance and the overhead in the financial plan.'),
    ('margin_calc', 'Event contribution margin (calculated)', '0.0%', 'IF({ingresos}=0,0,{profit}/{ingresos})',
     'It comes from the costs; it is never typed in.'),
    ('profit_sc', 'Event contribution without the commission ($)', USD, '{base}-{c_total}',
     'What the event would leave if the caterer paid no commission.'),
    ('can1', 'Amount kept if the client cancels: tier 1', USD, '{total}*{t_cn1}',
     'Furthest from the event (see the tier table). If the venue cancels: it refunds everything it collected.'),
    ('can2', 'Amount kept if the client cancels: tier 2', USD, '{total}*{t_cn2}', ''),
    ('can3', 'Amount kept if the client cancels: tier 3', USD, '{total}*{t_cn3}', ''),
    ('can4', 'Amount kept if the client cancels: tier 4', USD, '{total}*{t_cn4}', ''),
    ('can5', 'Amount kept if the client cancels: tier 5', USD, '{total}*{t_cn5}',
     'Last tier, up to the event day. The agreement credits back what you recover by rebooking the date and the costs you '
     'save by not holding the event.'),
]
S = OrderedDict((k, F_CALC + 1 + i) for i, (k, *_r) in enumerate(SALIDAS))

# --- bloques del pie (tablas) ---------------------------------------------------------------------------------------
F_TIPOS = F_CALC + 1 + len(SALIDAS) + 1
F_TEMP = F_TIPOS + 1 + len(TIPOS) + 1
F_DIA = F_TEMP + 1 + len(TEMPORADAS) + 1
F_MES = F_DIA + 1 + len(DIAS) + 1
F_SEM = F_MES + 1 + len(MESES) + 1
F_EXT = F_SEM + 1 + len(SEMANA) + 1
X = OrderedDict((k, F_EXT + 1 + i) for i, k in enumerate(EXTRAS))
F_CAN = F_EXT + 1 + len(EXTRAS) + 1
CT = OrderedDict((k, F_CAN + 1 + i) for i, k in enumerate(CANAL_TAB))
F_TRA = F_CAN + 1 + len(CANAL_TAB) + 1

_BASE = dict(h_aux=16, wage=18.0, ss=0.10, cost_ext=0.40, redondeo=50, canon=0.10, iva=0.08, iva_menu=0.08)
EJEMPLOS = OrderedDict([
    ('Saturday wedding in peak season', dict(
        _BASE, canal='Couple', tipo='Wedding', fecha=serial(datetime.date(2027, 10, 9)), guests=130, n_cer=1, n_hex=2,
        n_mon=1, n_reh=0)),
    ('Friday wedding in shoulder season', dict(
        _BASE, canal='Couple', tipo='Wedding', fecha=serial(datetime.date(2027, 7, 16)), guests=100, n_cer=1, n_hex=0,
        n_mon=0, n_reh=1, h_aux=14)),
    ('Weekday corporate event', dict(
        _BASE, canal='Corporate', tipo='Corporate event', fecha=serial(datetime.date(2027, 4, 21)), guests=50, n_cer=0,
        n_hex=1, n_mon=0, n_reh=0, h_aux=10)),
])
EJ_NOTAS = OrderedDict([
    ('Saturday wedding in peak season', 'Couple, 130 guests, a Saturday in October: peak season (+15%) and the reference '
     'day (Saturday, no adjustment). Ceremony on site, two additional hours and the premium decor package. The couple sees '
     'the price with sales tax included. The commission is what the exclusive caterer would pay the venue at 10% (an '
     'assumption): the couple does not pay it. Three payments: retainer, second installment six months before and final '
     'balance 30 days before.'),
    ('Friday wedding in shoulder season', 'Couple, 100 guests, a Friday in July: shoulder season (no adjustment) and '
     'Friday (−15%). Ceremony on site and a rehearsal the day before. Same payment schedule and cancellation tiers as the '
     'first example: only the venue rate changes.'),
    ('Weekday corporate event', 'Corporate client, 50 guests, a Wednesday in April: shoulder season and a weekday (−30%). '
     'The 50 guests are below the season minimum (75): the quote uses 75 and the commission is estimated on them. Quoted '
     '«plus sales tax» and paid on another schedule (30% to book, 30% two months before and 40% 15 days before) with other '
     'cancellation tiers.'),
])


# ---------------------------------------------------------------------------------------------------------------------
# Cálculo en Python (espejo de las fórmulas) — la fuente del autotest
# ---------------------------------------------------------------------------------------------------------------------
def _r(x):
    """ROUND de Excel (la mitad se aleja de cero)."""
    return math.floor(x + 0.5) if x >= 0 else -math.floor(-x + 0.5)


def _fecha(n):
    return EPOCA + datetime.timedelta(days=int(n))


def calcular(e):
    r = OrderedDict()
    ic = CANALES.index(e['canal'])
    t = lambda k: CANAL_TAB[k][1][ic]
    d = _fecha(e['fecha'])
    r['temporada'] = MESES[d.month - 1][1]
    r['dia'] = SEMANA[d.isoweekday() % 7][1]
    r['aj_t'] = TEMPORADAS[r['temporada']][0]
    r['aj_d'] = DIAS[r['dia']]
    tarifa, menu, cfijo = TIPOS[e['tipo']]
    r['tarifa'] = tarifa
    r['espacio'] = _r(tarifa * (1 + r['aj_t']) * (1 + r['aj_d']) / e['redondeo']) * e['redondeo']
    r['min_inv'] = TEMPORADAS[r['temporada']][1]
    r['garant'] = max(e['guests'], r['min_inv'])
    r['aviso'] = ('Below the minimum: the quote uses the season minimum' if e['guests'] < r['min_inv']
                  else 'Meets the minimum')
    for k in EXTRAS:
        r['e_' + k] = e['n_' + k] * EXTRAS[k][1]
    r['extras'] = r['e_cer'] + r['e_hex'] + r['e_mon'] + r['e_reh']
    r['base'] = r['espacio'] + r['extras']
    r['iva_amt'] = r['base'] * e['iva']
    r['total'] = r['base'] + r['iva_amt']
    r['quoted'] = r['total'] if e['canal'] == 'Couple' else r['base']
    r['eur_inv'] = r['base'] / r['garant']
    r['cb1'] = r['total'] * t('pc_1')
    r['cb2'] = r['total'] * t('pc_2')
    r['f_2'] = e['fecha'] - t('d_2')
    r['cb3'] = r['total'] * t('pc_3')
    r['f_3'] = e['fecha'] - t('d_3')
    r['menu_tot'] = menu * r['garant']
    r['canon_est'] = r['menu_tot'] * e['canon']
    r['menu_cli'] = r['menu_tot'] * (1 + e['iva_menu'])
    r['presup'] = r['total'] + r['menu_cli']
    r['c_fijo'] = cfijo
    r['c_aux'] = e['h_aux'] * e['wage'] * (1 + e['ss'])
    r['c_ext'] = r['extras'] * e['cost_ext']
    r['c_total'] = r['c_fijo'] + r['c_aux'] + r['c_ext']
    r['ingresos'] = r['base'] + r['canon_est']
    r['profit'] = r['ingresos'] - r['c_total']
    r['margin_calc'] = 0 if r['ingresos'] == 0 else r['profit'] / r['ingresos']
    r['profit_sc'] = r['base'] - r['c_total']
    for i in (1, 2, 3, 4, 5):
        r['can%d' % i] = r['total'] * t('cn%d' % i)
    return r


# ---------------------------------------------------------------------------------------------------------------------
# Fórmulas
# ---------------------------------------------------------------------------------------------------------------------
def _rng(col, f1, f2):
    return "'%s'!$%s$%d:$%s$%d" % (H1, col, f1, col, f2)


def _busca(col, f0, n, clave):
    return 'INDEX(%s,MATCH(%s,%s,0))' % (_rng(col, f0 + 1, f0 + n), clave, _rng('A', f0 + 1, f0 + n))


def refs(col):
    """Referencias a celdas de la columna `col` (B en Calculator; B/C/D en Examples) + búsquedas en las tablas."""
    ref = dict((k, '$%s$%d' % (col, f)) for k, f in K.items())
    ref.update((k, '$%s$%d' % (col, f)) for k, f in S.items())
    ref['t_tarifa'] = _busca('B', F_TIPOS, len(TIPOS), ref['tipo'])
    ref['t_menu'] = _busca('C', F_TIPOS, len(TIPOS), ref['tipo'])
    ref['t_cfijo'] = _busca('D', F_TIPOS, len(TIPOS), ref['tipo'])
    ref['t_aj_t'] = _busca('B', F_TEMP, len(TEMPORADAS), ref['temporada'])
    ref['t_min'] = _busca('C', F_TEMP, len(TEMPORADAS), ref['temporada'])
    ref['t_aj_d'] = _busca('B', F_DIA, len(DIAS), ref['dia'])
    ref['t_temporada'] = 'INDEX(%s,MONTH(%s))' % (_rng('B', F_MES + 1, F_MES + len(MESES)), ref['fecha'])
    ref['t_dia'] = 'INDEX(%s,WEEKDAY(%s))' % (_rng('B', F_SEM + 1, F_SEM + len(SEMANA)), ref['fecha'])
    for k, f in X.items():
        ref['x_' + k] = "'%s'!$B$%d" % (H1, f)
    cab = "'%s'!$B$%d:$C$%d" % (H1, F_CAN, F_CAN)
    for k, f in CT.items():
        ref['t_' + k] = "INDEX('%s'!$B$%d:$C$%d,MATCH(%s,%s,0))" % (H1, f, f, ref['canal'], cab)
    return ref


def a_excel(expr, ref):
    return re.sub(r'\{(\w+)\}', lambda m: ref[m.group(1)], expr)


def _f(ws, coord, v, b=False, sz=11, i=False, fmt=None, wrap=False, fill=None, locked=True):
    c = ws[coord]
    c.value = v
    c.font = Font(name=FUENTE, size=sz, bold=b, italic=i)
    if fmt:
        c.number_format = fmt
    c.alignment = Alignment(wrap_text=wrap, vertical='top')
    if fill:
        c.fill = fill
    c.protection = Protection(locked=locked)
    return c


def _nota(ws, fila, texto, desde='C'):
    """Nota combinada en C:F (ancho útil ≈ 110 caracteres)."""
    if not texto:
        return
    _f(ws, '%s%d' % (desde, fila), texto, wrap=True)
    ws.merge_cells('%s%d:F%d' % (desde, fila, fila))
    ws.row_dimensions[fila].height = max(16, 15 * math.ceil(len(texto) / (105.0 if desde == 'C' else 80.0)))


def _imprimir(ws, horizontal=False):
    ws.page_setup.paperSize = M.PAPEL_EN                                         # 1 = US Letter
    ws.page_setup.orientation = 'landscape' if horizontal else 'portrait'
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.oddFooter.center.text = FOOTER
    ws.print_options.horizontalCentered = True
    ws.protection.sheet = True                                                    # sin contraseña
    ws.sheet_view.showGridLines = False


def _dv(ws, tipo, celdas, f1=None, f2=None, op=None, titulo='Invalid value', msg=''):
    dv = DataValidation(type=tipo, formula1=f1, formula2=f2, operator=op, allow_blank=False, showErrorMessage=True,
                        errorTitle=titulo, error=msg)
    ws.add_data_validation(dv)
    for c in celdas:
        dv.add(c)


def _cab(ws, f, titulo, cols):
    _f(ws, 'A%d' % f, titulo, b=True, fill=GRIS)
    for col, t in cols:
        _f(ws, '%s%d' % (col, f), t, b=True, fill=GRIS, wrap=True)
    ws.row_dimensions[f].height = 32


def hoja_calculadora(wb):
    ws = wb.active
    ws.title = H1
    for col, w in (('A', 66), ('B', 22), ('C', 17), ('D', 17), ('E', 17), ('F', 56)):
        ws.column_dimensions[col].width = w
    _f(ws, 'A1', 'WEDDING VENUE PRICING & QUOTE CALCULATOR', b=True, sz=14)
    _f(ws, 'A2', 'Type only in the GREEN cells: everything else is calculated. Amounts exclude sales tax unless the row '
                 'says otherwise. The default values are Kit estimates: replace them with your own venue\'s.', wrap=True)
    ws.merge_cells('A2:F2')
    ws.row_dimensions[2].height = 32
    for col, t in zip('ABC', ('EVENT DETAILS', 'Value', 'Notes')):
        _f(ws, '%s4' % col, t, b=True, fill=GRIS)
    for k, rot, fmt, val, nota in ENTRADAS:
        f = K[k]
        _f(ws, 'A%d' % f, rot)
        _f(ws, 'B%d' % f, val, fmt=fmt if fmt != '@' else None, fill=VERDE, locked=False)
        _nota(ws, f, nota)
    for col, t in zip('ABC', ('CALCULATION', 'Value', 'Notes')):
        _f(ws, '%s%d' % (col, F_CALC), t, b=True, fill=GRIS)
    ref = refs('B')
    for k, rot, fmt, expr, nota in SALIDAS:
        f = S[k]
        neg = k in ('espacio', 'total', 'quoted', 'canon_est', 'margin_calc', 'presup')
        _f(ws, 'A%d' % f, rot, b=neg)
        _f(ws, 'B%d' % f, '=' + a_excel(expr, ref), fmt=fmt, b=neg)
        _nota(ws, f, nota)
    _cab(ws, F_TIPOS, 'EVENT TYPE → VENUE RATE, CATERER\'S FOOD & BEVERAGE AND DIRECT COST (excl. sales tax)',
         [('B', 'Venue rate (Saturday, shoulder season)'), ('C', 'Caterer\'s food & beverage per guest'),
          ('D', 'Direct event cost')])
    for i, (nombre, (tarifa, menu, cfijo)) in enumerate(TIPOS.items(), 1):
        _f(ws, 'A%d' % (F_TIPOS + i), nombre)
        _f(ws, 'B%d' % (F_TIPOS + i), tarifa, fmt=USD, fill=VERDE, locked=False)
        _f(ws, 'C%d' % (F_TIPOS + i), menu, fmt=USD, fill=VERDE, locked=False)
        _f(ws, 'D%d' % (F_TIPOS + i), cfijo, fmt=USD, fill=VERDE, locked=False)
    _f(ws, 'E%d' % (F_TIPOS + 1), 'Kit estimates (published barn venue rate cards range roughly $3,500-$10,000 for a peak '
                                  'Saturday: get local quotes). Direct cost: cleaning, restrooms, generator fuel, parking '
                                  'attendants and security.', i=True, wrap=True)
    ws.merge_cells('E%d:F%d' % (F_TIPOS + 1, F_TIPOS + 4))
    _cab(ws, F_TEMP, 'SEASON → RATE ADJUSTMENT AND GUEST MINIMUM',
         [('B', 'Adjustment (+ premium / − discount)'), ('C', 'Guest minimum')])
    for i, (nombre, (aj, mn)) in enumerate(TEMPORADAS.items(), 1):
        _f(ws, 'A%d' % (F_TEMP + i), nombre)
        _f(ws, 'B%d' % (F_TEMP + i), aj, fmt=PCT, fill=VERDE, locked=False)
        _f(ws, 'C%d' % (F_TEMP + i), mn, fmt=ENT, fill=VERDE, locked=False)
    _f(ws, 'D%d' % (F_TEMP + 1), 'Kit example. Peak season carries a premium and off-peak a discount; the event type rates '
                                 'are for a Saturday in shoulder season.', i=True, wrap=True)
    ws.merge_cells('D%d:F%d' % (F_TEMP + 1, F_TEMP + 3))
    _cab(ws, F_DIA, 'TYPE OF DAY → RATE ADJUSTMENT', [('B', 'Adjustment (+ premium / − discount)')])
    for i, (nombre, aj) in enumerate(DIAS.items(), 1):
        _f(ws, 'A%d' % (F_DIA + i), nombre)
        _f(ws, 'B%d' % (F_DIA + i), aj, fmt=PCT, fill=VERDE, locked=False)
    _f(ws, 'C%d' % (F_DIA + 1), 'Kit example: venue rate cards usually discount Fridays, Sundays and weekdays. The season '
                                'and day adjustments multiply.', i=True, wrap=True)
    ws.merge_cells('C%d:F%d' % (F_DIA + 1, F_DIA + 3))
    _cab(ws, F_MES, 'MONTH → SEASON', [('B', 'Season')])
    for i, (mes, temp) in enumerate(MESES, 1):
        _f(ws, 'A%d' % (F_MES + i), mes)
        _f(ws, 'B%d' % (F_MES + i), temp, fill=VERDE, locked=False)
    _f(ws, 'C%d' % (F_MES + 1), 'Fall is the busiest wedding season in the US, with October, September and June as the '
                                'peaks (Kit pattern). Adjust the months to your area.', i=True, wrap=True)
    ws.merge_cells('C%d:F%d' % (F_MES + 1, F_MES + 3))
    _cab(ws, F_SEM, 'WEEKDAY → TYPE OF DAY', [('B', 'Type of day')])
    for i, (dia, tipo) in enumerate(SEMANA, 1):
        _f(ws, 'A%d' % (F_SEM + i), dia)
        _f(ws, 'B%d' % (F_SEM + i), tipo, fill=VERDE, locked=False)
    _cab(ws, F_EXT, 'VENUE EXTRA SERVICES (price excl. sales tax)', [('B', 'Price')])
    for k, (rot, p) in EXTRAS.items():
        _f(ws, 'A%d' % X[k], rot)
        _f(ws, 'B%d' % X[k], p, fmt=USD, fill=VERDE, locked=False)
    _f(ws, 'C%d' % (F_EXT + 1), 'Kit example prices: type your own.', i=True, wrap=True)
    ws.merge_cells('C%d:F%d' % (F_EXT + 1, F_EXT + 2))
    _cab(ws, F_CAN, 'POLICY BY CLIENT', [('B', CANALES[0]), ('C', CANALES[1])])
    for k, (rot, vals, nota) in CANAL_TAB.items():
        f = CT[k]
        _f(ws, 'A%d' % f, rot)
        for col, v in zip('BC', vals):
            _f(ws, '%s%d' % (col, f), v, fmt=ENT if k.startswith('d_') else PCT, fill=VERDE, locked=False)
        _f(ws, 'D%d' % f, nota, wrap=True)
        ws.merge_cells('D%d:F%d' % (f, f))
        ws.row_dimensions[f].height = max(16, 15 * math.ceil(len(nota) / 78.0))
    _cab(ws, F_TRA, 'CANCELLATION TIERS (months before the event)',
         [('B', 'From (months, inclusive)'), ('C', 'Up to (months, exclusive)')])
    for i, m in enumerate(TRAMOS):
        f = F_TRA + 1 + i
        _f(ws, 'A%d' % f, 'Tier %d' % (i + 1))
        _f(ws, 'B%d' % f, m, fmt=ENT, fill=VERDE, locked=False)
        _f(ws, 'C%d' % f, 'no limit' if i == 0 else '=$B$%d' % (f - 1), fmt='General' if i == 0 else ENT)
    f = F_TRA + 1 + len(TRAMOS)
    _f(ws, 'A%d' % f, 'No gaps: each tier starts where the previous one ends and the last one runs to the event day. The '
                      'amounts kept rise as the date gets closer and are meant as a reasonable estimate of the venue\'s '
                      'loss, not a penalty: the agreement credits back what you recover by rebooking the date. If the venue '
                      'cancels, it refunds everything it collected. The percentages and tiers are Kit examples: set them '
                      'to your real cost of losing the date and have an attorney in your state review them.',
       i=True, wrap=True)
    ws.merge_cells('A%d:F%d' % (f, f))
    ws.row_dimensions[f].height = 62
    ref = refs('B')
    cel = lambda k: ref[k].replace('$', '')
    _dv(ws, 'list', [cel('canal')], f1='"%s"' % ','.join(CANALES), titulo='Invalid client', msg='Choose Couple or Corporate.')
    _dv(ws, 'list', [cel('tipo')], f1='=$A$%d:$A$%d' % (F_TIPOS + 1, F_TIPOS + len(TIPOS)), titulo='Invalid event type',
        msg='Choose an event type from the list.')
    _dv(ws, 'date', [cel('fecha')], f1=str(serial(datetime.date(2026, 1, 1))), op='greaterThanOrEqual',
        titulo='Invalid date', msg='Type a full date from 2026 on (mm/dd/yyyy).')
    _dv(ws, 'whole', [cel('guests')], f1='1', f2='100000', op='between', titulo='Invalid number',
        msg='Type a whole number of 1 or more.')
    _dv(ws, 'whole', [cel(k) for k in ('n_cer', 'n_hex', 'n_mon', 'n_reh')]
        + ['C%d' % (F_TEMP + i) for i in range(1, len(TEMPORADAS) + 1)]
        + ['B%d' % (F_TRA + 1 + i) for i in range(len(TRAMOS))] + ['%s%d' % (c, CT[k]) for k in ('d_2', 'd_3') for c in 'BC'],
        f1='0', f2='100000', op='between', titulo='Invalid number', msg='Type a whole number of 0 or more.')
    _dv(ws, 'decimal', [cel(k) for k in ('h_aux', 'wage')] + ['B%d' % X[k] for k in X]
        + ['%s%d' % (c, F_TIPOS + i) for c in 'BCD' for i in range(1, len(TIPOS) + 1)],
        f1='0', op='greaterThanOrEqual', titulo='Invalid amount', msg='Type a number greater than or equal to 0.')
    _dv(ws, 'decimal', [cel(k) for k in ('ss', 'cost_ext', 'canon', 'iva', 'iva_menu')]
        + ['%s%d' % (c, CT[k]) for k in CANAL_TAB if not k.startswith('d_') for c in 'BC'],
        f1='0', f2='1', op='between', titulo='Invalid percentage', msg='Type a percentage between 0 and 1 (0.08 = 8%).')
    _dv(ws, 'decimal', ['B%d' % (F_TEMP + i) for i in range(1, len(TEMPORADAS) + 1)]
        + ['B%d' % (F_DIA + i) for i in range(1, len(DIAS) + 1)], f1='-0.9', f2='2', op='between',
        titulo='Invalid adjustment', msg='Type an adjustment between −90% and +200% (−0.20 = −20%).')
    _dv(ws, 'list', ['B%d' % (F_MES + i) for i in range(1, len(MESES) + 1)], f1='"%s"' % ','.join(TEMPORADAS),
        titulo='Invalid season', msg='Choose Peak, Shoulder or Off-peak.')
    _dv(ws, 'list', ['B%d' % (F_SEM + i) for i in range(1, len(SEMANA) + 1)], f1='"%s"' % ','.join(DIAS),
        titulo='Invalid type of day', msg='Choose Saturday, Friday or Sunday, or Weekday.')
    _dv(ws, 'whole', [cel('redondeo')], f1='1', f2='100000', op='between', titulo='Invalid rounding',
        msg='Type a multiple of $1 or more (the fee is divided by it).')
    ws.freeze_panes = 'A5'
    _imprimir(ws)


def hoja_ejemplos(wb):
    ws = wb.create_sheet(H2)
    ws.column_dimensions['A'].width = 66
    for col in 'BCD':
        ws.column_dimensions[col].width = 28
    _f(ws, 'A1', '3 EXAMPLE QUOTES — with the same formulas as the "%s" tab' % H1, b=True, sz=14)
    _f(ws, 'A2', 'Kit estimates to show the calculation, not market prices: change the inputs in "%s" to quote your own '
                 'events. The tables (rates, seasons, days, extras, payments and cancellation) are at the end of that tab '
                 'and also drive these examples.' % H1, wrap=True)
    ws.merge_cells('A2:D2')
    ws.row_dimensions[2].height = 32
    _f(ws, 'A4', 'EVENT DETAILS', b=True, fill=GRIS)
    for j, nombre in enumerate(EJEMPLOS):
        _f(ws, '%s4' % 'BCD'[j], nombre, b=True, fill=GRIS, wrap=True)
    ws.row_dimensions[4].height = 48
    for k, rot, fmt, val, nota in ENTRADAS:
        _f(ws, 'A%d' % K[k], rot)
        for j, e in enumerate(EJEMPLOS.values()):
            _f(ws, '%s%d' % ('BCD'[j], K[k]), e[k], fmt=fmt if fmt != '@' else None, wrap=(fmt == '@'))
    _f(ws, 'A%d' % F_CALC, 'RESULT', b=True, fill=GRIS)
    for j in range(3):
        _f(ws, '%s%d' % ('BCD'[j], F_CALC), list(EJEMPLOS)[j], b=True, fill=GRIS, wrap=True)
    ws.row_dimensions[F_CALC].height = 48
    for k, rot, fmt, expr, nota in SALIDAS:
        neg = k in ('espacio', 'total', 'quoted', 'canon_est', 'margin_calc', 'presup')
        _f(ws, 'A%d' % S[k], rot, b=neg)
        for j in range(3):
            col = 'BCD'[j]
            _f(ws, '%s%d' % (col, S[k]), '=' + a_excel(expr, refs(col)), fmt=fmt, b=neg, wrap=(fmt == 'General'))
    f = S[SALIDAS[-1][0]] + 2
    _f(ws, 'A%d' % f, 'WHAT EACH EXAMPLE SHOWS', b=True, fill=GRIS)
    for j, (nombre, nota) in enumerate(EJ_NOTAS.items()):
        _f(ws, 'A%d' % (f + 1 + j), nombre, b=True, wrap=True)
        _f(ws, 'B%d' % (f + 1 + j), nota, wrap=True)
        ws.merge_cells('B%d:D%d' % (f + 1 + j, f + 1 + j))
        ws.row_dimensions[f + 1 + j].height = 15 * max(3, math.ceil(len(nota) / 80.0))
    ws.freeze_panes = 'B5'
    _imprimir(ws, horizontal=True)


COOLING = ('If you sign couples at bridal shows or off-site, ask your attorney whether the FTC Cooling-Off Rule or your '
           'state\'s home-solicitation law applies; many venues give a short cancellation window for show-floor signings '
           'anyway.')
INSTRUCCIONES = [
    ('HOW TO USE — Wedding Venue Pricing & Quote Calculator', True),
    ('It turns your rate card into a full quote for a specific event: the season and type of day of the date, the guest '
     'minimum, extra services, sales tax, total, payment schedule, the estimated caterer commission, the event '
     'contribution and what you keep if the client cancels. Type only in the GREEN cells; the margin is calculated.', False),
    ('1. Choose the client (Couple or Corporate), the event type, the full date and the expected guests.', False),
    ('2. The date sets the season (month table) and the type of day (weekday table). Each one carries a premium or a '
     'discount in its own visible table; both multiply the event type rate.', False),
    ('3. Each season has a guest minimum: the quote uses the higher of the expected guests and the minimum, and the '
     'commission is estimated on that guaranteed guest count.', False),
    ('4. Type how many extra services the event takes (ceremony on site, additional hours, premium decor package, '
     'rehearsal the day before). Their prices are in the "Venue extra services" table and are Kit examples: type your '
     'own.', False),
    ('5. Type the venue staff hours, the hourly wage, the payroll tax rate and the direct cost of the extras. The direct '
     'event cost (cleaning, restrooms, generator fuel, parking and security) is in the event type table.', False),
    ('SALES TAX, COMMISSION AND THE CATERER\'S MENU', True),
    ('The venue bills its rental fee and services at the sales tax rate you type (8% in the example). Whether a venue '
     'rental fee is taxable depends on the state (taxable in some, exempt in others, sometimes only when it is billed with '
     'catering): ask your state\'s department of revenue. The exclusive caterer bills food and the bar with its own sales '
     'tax, service charge and gratuity; they are not venue revenue, and the calculator shows them only as a reference for '
     'the client.', False),
    ('The caterer commission is what the exclusive caterer pays the venue for the exclusivity. Here it is an editable '
     'assumption (10% of the caterer\'s food & beverage sales, net of sales tax, service charge and gratuities): it is not '
     'a market figure, negotiate it. The caterer pays it, not the client; the plan treats it as a fee between businesses '
     '(0% sales tax in the example): confirm it with your CPA.', False),
    ('The event contribution does not deduct rent, salaried staff, utilities, insurance or overhead. Those are in the '
     'financial plan.', False),
    ('PAYMENTS AND CANCELLATION', True),
    ('Three payments: the retainer at signing, a second installment and the final balance before the event, each with its '
     'due date. The retainer is applied to the price. Issue a receipt for every payment. Couple: 25%, 35% six months '
     'before and 40% 30 days before. Corporate: 30%, 30% two months before and 40% 15 days before. These are Kit '
     'examples.', False),
    ('Cancellation is written in tiers by months, with no gaps: the first starts at 12 months or more and the last runs to '
     'the event day. The amounts kept rise as the date gets closer and should be a reasonable estimate of the venue\'s '
     'loss, not a penalty: the agreement credits back what you recover by rebooking the date. Never write '
     '"non-refundable" without the schedule. If the venue cancels, it refunds everything it collected. Have an attorney in '
     'your state review the schedule (the kit\'s Venue Rental Agreement Templates use the same tiers).', False),
    (COOLING, False),
    ('KIT ESTIMATES', True),
    ('The default values and the three examples are estimates to show the calculation, not market prices. Replace them '
     'with your rates, extras, costs and tiers before you quote.', False),
]


def hoja_instrucciones(wb):
    ws = wb.create_sheet(H3)
    ws.column_dimensions['A'].width = 110
    f = 1
    for t, b in INSTRUCCIONES:
        _f(ws, 'A%d' % f, t, b=b, sz=13 if f == 1 else 11, wrap=True)
        if not b:
            ws.row_dimensions[f].height = max(16, 15 * math.ceil(len(t) / 120.0))
        f += 2 if b and f > 1 else 1
    f += 1
    for t in ('To edit a cell that is not green: Review → Unprotect Sheet (no password).',
              'More templates and kits in the catalog at ' + M.URL_TIENDA, M.FIRMA_EN, VERSION_LINEA, M.AVISO_EN):
        _f(ws, 'A%d' % f, t, wrap=True, i=t == M.AVISO_EN)
        f += 1
    _imprimir(ws)


def docprops(wb):
    pr = wb.properties
    pr.title = '%s · %s' % (TITULO, PRODUCTO)
    pr.subject = '%s · v%s' % (PRODUCTO, VERSION)
    pr.keywords = KEYWORDS
    pr.description = URL
    pr.category = 'AI Chef Pro · Digital products'
    pr.creator = pr.lastModifiedBy = 'AI Chef Pro'


def generar(destino):
    wb = openpyxl.Workbook()
    hoja_calculadora(wb)
    hoja_ejemplos(wb)
    hoja_instrucciones(wb)
    docprops(wb)
    os.makedirs(os.path.dirname(destino), exist_ok=True)
    wb.save(destino)
    import inject_cache as IC                                                      # noqa: E402
    nf, inj, fail = IC.inject(destino)
    print('%s: formulas=%d cache_inyectado=%d fallos_pycel=%d' % (os.path.basename(destino), nf, inj, fail))
    return fail


def _num(x):
    """Valor leído de la caché → número (las fechas vuelven como datetime)."""
    if isinstance(x, datetime.datetime):
        return serial(x.date())
    if isinstance(x, datetime.date):
        return serial(x)
    return x


def _distinto(x, v):
    x = _num(x)
    if isinstance(v, str):
        return x != v
    return not isinstance(x, (int, float)) or abs(float(x) - v) > 0.005


def autotest(destino):
    from pycel import ExcelCompiler
    err, n = [], 0
    xl = ExcelCompiler(filename=destino)
    wv = openpyxl.load_workbook(destino, data_only=True)
    wf = openpyxl.load_workbook(destino)
    casos = [(H2, 'BCD'[j], e, nombre) for j, (nombre, e) in enumerate(EJEMPLOS.items())]
    casos.append((H1, 'B', OrderedDict((k, v) for k, _r2, _f2, v, _n in ENTRADAS), 'main tab'))
    for hoja, col, e, nombre in casos:
        esperado = calcular(e)
        for k, v in esperado.items():
            c = '%s%d' % (col, S[k])
            n += 1
            for lab, x in (('pycel', xl.evaluate("'%s'!%s" % (hoja, c))), ('caché', wv[hoja][c].value)):
                if _distinto(x, v):
                    err.append('%s · %s %s (%s): %s %r ≠ esperado %r' % (nombre, hoja, c, k, lab, x, v))
        if not 20 <= esperado['eur_inv'] <= 200:
            err.append('%s: venue rental & services $%.2f per guest outside $20-200' % (nombre, esperado['eur_inv']))
    pe = [calcular(e) for e in EJEMPLOS.values()]
    if (pe[0]['temporada'], pe[0]['dia']) != ('Peak', 'Saturday') or \
            (pe[1]['temporada'], pe[1]['dia']) != ('Shoulder', 'Friday or Sunday') or \
            (pe[2]['temporada'], pe[2]['dia']) != ('Shoulder', 'Weekday') or \
            pe[2]['garant'] <= EJEMPLOS['Weekday corporate event']['guests']:
        err.append('los ejemplos no son sábado peak / viernes shoulder / weekday con mínimo aplicado')
    base = OrderedDict((k, v) for k, _r2, _f2, v, _n in ENTRADAS)
    propios = [('own 1: Corporate, Monday in January, 20 guests, micro wedding',
                dict(base, canal='Corporate', tipo='Micro wedding / elopement', fecha=serial(datetime.date(2027, 1, 18)),
                     guests=20, n_cer=0, n_hex=0, n_mon=0, h_aux=4, redondeo=10)),
               ('own 2: Couple, Sunday in September, 150 guests, private celebration',
                dict(base, canal='Couple', tipo='Private celebration', fecha=serial(datetime.date(2027, 9, 12)), guests=150,
                     n_cer=2, n_hex=3, n_mon=2, n_reh=1, canon=0.12, h_aux=20)),
               ('own 3: Corporate, Saturday in August, 85 guests, corporate event',
                dict(base, canal='Corporate', tipo='Corporate event', fecha=serial(datetime.date(2027, 8, 7)), guests=85,
                     n_cer=0, n_hex=4, n_mon=1, n_reh=0, iva=0.0625, iva_menu=0.0625))]
    for nombre, e in propios:
        esperado = calcular(e)
        xl2 = ExcelCompiler(filename=destino)
        for k in S:
            xl2.evaluate("'%s'!B%d" % (H1, S[k]))
        for k, v in e.items():
            xl2.set_value("'%s'!B%d" % (H1, K[k]), v)
        for k, v in esperado.items():
            n += 1
            x = xl2.evaluate("'%s'!B%d" % (H1, S[k]))
            if _distinto(x, v):
                err.append('%s: %s pycel %r ≠ esperado %r' % (nombre, k, x, v))
    e = dict(EJEMPLOS['Saturday wedding in peak season'])
    cs = [calcular(dict(e, canal=c)) for c in CANALES]
    if cs[0]['cb1'] == cs[1]['cb1'] or cs[0]['can2'] == cs[1]['can2'] or cs[0]['quoted'] == cs[1]['quoted']:
        err.append('el selector de cliente no cambia el cobro, la cancelación o el importe presupuestado')
    for ic, canal in enumerate(CANALES):
        s = sum(CANAL_TAB[k][1][ic] for k in ('pc_1', 'pc_2', 'pc_3'))
        if abs(s - 1.0) > 1e-9:
            err.append('cobro %s suma %.2f ≠ 100 %%' % (canal, s))
        v = [CANAL_TAB['cn%d' % i][1][ic] for i in range(1, 6)]
        if v != sorted(v) or v[-1] != 1.0:
            err.append('cancelación %s: tramos no crecientes o el último no es el 100 %% %r' % (canal, v))
        if v[0] > CANAL_TAB['pc_1'][1][ic]:
            err.append('cancelación %s: el tramo 1 retiene más que el retainer' % canal)
        if CANAL_TAB['d_2'][1][ic] <= CANAL_TAB['d_3'][1][ic]:
            err.append('cobro %s: el segundo pago no es anterior a la liquidación' % canal)
    if TRAMOS != sorted(TRAMOS, reverse=True) or len(set(TRAMOS)) != len(TRAMOS) or TRAMOS[-1] != 0 or len(TRAMOS) != 5:
        err.append('tramos de cancelación con huecos o desordenados: %r' % TRAMOS)
    ws = wf[H1]
    for i in range(len(TRAMOS)):
        v = ws['C%d' % (F_TRA + 1 + i)].value
        if i and v != '=$B$%d' % (F_TRA + i):
            err.append('tramo %d: «hasta» no enlaza con el «desde» anterior (%r)' % (i + 1, v))
    if any(t not in TEMPORADAS for _m, t in MESES) or any(d not in DIAS for _s, d in SEMANA) or len(MESES) != 12 \
            or len(SEMANA) != 7:
        err.append('tabla de meses o de días de la semana incompleta')
    if [m for m, t in MESES if t == 'Peak'] != ['May', 'June', 'September', 'October']:
        err.append('meses peak distintos de D168')
    gf = motor.gate_formatos(wf, FICHERO)
    if gf:
        err.append('gate_formatos: %d celdas, p. ej. %r' % (len(gf), gf[:3]))
    dvs = dict((str(rng), dv) for dv in ws.data_validations.dataValidation for rng in dv.sqref.ranges)
    c = 'B%d' % K['canal']
    if dvs.get(c) is None or dvs[c].type != 'list' or 'Couple,Corporate' not in (dvs[c].formula1 or ''):
        err.append('DV de lista incorrecta en %s' % c)
    c = 'B%d' % K['tipo']
    if dvs.get(c) is None or dvs[c].type != 'list' or not (dvs[c].formula1 or '').startswith('=$A$'):
        err.append('DV de lista del rango incorrecta en %s' % c)
    for k, tp in (('fecha', 'date'), ('guests', 'whole'), ('n_cer', 'whole')):
        dv = dvs.get('B%d' % K[k])
        if dv is None or dv.type != tp:
            err.append('DV %s ausente en %s' % (tp, k))
    textos = []
    for w in wf.worksheets:
        if w.page_setup.paperSize != 1 or not w.sheet_properties.pageSetUpPr.fitToPage or not w.protection.sheet \
                or w.oddFooter.center.text != FOOTER:
            err.append('%s: Letter, ajuste, protección o pie' % w.title)
        textos.append((w.title, '(pestaña)', w.title))
        for row in w.iter_rows():
            for c in row:
                v = c.value
                if isinstance(v, str) and c.data_type != 'f':
                    textos.append((w.title, c.coordinate, v))
                if c.data_type == 'f':
                    if wv[w.title][c.coordinate].value is None:
                        err.append('%s!%s: fórmula sin caché' % (w.title, c.coordinate))
                    for lit in re.findall(r'"((?:[^"]|"")*)"', c.value):
                        if any(ch.isalpha() for ch in lit):
                            textos.append((w.title, c.coordinate + ' (fórmula)', lit))
                    sin_refs = re.sub(r"'[^']+'!|\$?[A-Z]{1,2}\$?\d+|\"(?:[^\"]|\"\")*\"", '', str(c.value))
                    for num in re.findall(r'(?<![A-Za-z])\d+\.?\d*', sin_refs):
                        if num not in ('0', '1'):
                            err.append('%s!%s: número tecleado %s en la fórmula %s' % (w.title, c.coordinate, num, c.value))
                if c.fill.fgColor.rgb in ('00E8F5E9', 'FFE8F5E9') and c.protection.locked:
                    err.append('%s!%s: verde bloqueada' % (w.title, c.coordinate))
                if isinstance(c.number_format, str) and ('€' in c.number_format or 'dd/mm' in c.number_format):
                    err.append('%s!%s: formato ES %r' % (w.title, c.coordinate, c.number_format))
        for dv in w.data_validations.dataValidation:
            for t in (dv.error, dv.errorTitle):
                if t:
                    textos.append((w.title, 'DV %s' % dv.sqref, t))
    for h, cc, t in textos:
        r = [x for x in M.restos_espanol(t) if not x[0].startswith('oficio')]
        if r:
            err.append('%s!%s: restos de español %s «%s»' % (h, cc, r[:1], t[:60]))
        if M.no_latinos(t):
            err.append('%s!%s: caracteres no latinos' % (h, cc))
        for motivo, frag in M.cifras_prohibidas(t):
            err.append('%s!%s: %s «%s»' % (h, cc, motivo, frag))
        if re.search(r'(?i)\bnon-?refundable\b', t) and 'schedule' not in t:
            err.append('%s!%s: «non-refundable» sin la tabla' % (h, cc))
        if re.search(r'(?i)\bvalue\b.*\$\d|\bsave \$|total value|\bwas \$|\d+\s?% off\b', t):
            err.append('%s!%s: ancla de precio' % (h, cc))
    todo = '\n'.join(t for _h, _c, t in textos)
    for frase in (COOLING, 'a reasonable estimate of the venue\'s loss, not a penalty', 'rebooking the date',
                  'ask your state\'s department of revenue'):
        if frase not in todo:
            err.append('falta el texto %r' % frase)
    if re.search(r'(?i)cooling-off rule (?:applies|gives)', todo):
        err.append('afirma que la Cooling-Off Rule aplica (decisión de orquestador 2)')
    if M.FIRMA_EN not in todo or VERSION_LINEA not in todo or M.AVISO_EN not in todo:
        err.append('falta la firma, la línea de versión o el aviso')
    if wf.sheetnames != [H1, H2, H3]:
        err.append('pestañas %r' % wf.sheetnames)
    p = wf.properties
    if p.creator != 'AI Chef Pro' or p.description != URL or 'AI Chef Pro' not in (p.keywords or '') \
            or not (p.title or '').startswith(TITULO):
        err.append('docProps D20: creator=%r title=%r description=%r' % (p.creator, p.title, p.description))
    for x in err[:80]:
        print('ERROR', x)
    print('autotest wedding venue pricing calculator: %d comprobaciones (pycel + caché en ejemplos y hoja principal, pycel '
          'en 3 casos propios) sobre %d textos · %s' % (n * 2, len(textos), 'VERDE' if not err else 'ROJO (%d errores)' % len(err)))
    for nombre, e in EJEMPLOS.items():
        r = calcular(e)
        print('  %-36s %s/%s · venue fee %9.2f · base %9.2f · total %9.2f · retainer %8.2f · commission %8.2f · margin %5.1f %%'
              % (nombre, r['temporada'], r['dia'], r['espacio'], r['base'], r['total'], r['cb1'], r['canon_est'],
                 r['margin_calc'] * 100))
    return 1 if err else 0


if __name__ == '__main__':
    ap = argparse.ArgumentParser()
    ap.add_argument('--salida', required=True, help='carpeta de salida')
    ap.add_argument('--autotest', action='store_true')
    a = ap.parse_args()
    destino = os.path.join(a.salida, FICHERO)
    if not a.autotest and generar(destino):
        sys.exit('ERROR: fallos de pycel en inject_cache')
    sys.exit(autotest(destino))
