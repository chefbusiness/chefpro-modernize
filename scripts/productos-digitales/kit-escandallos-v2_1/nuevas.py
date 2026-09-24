#!/usr/bin/env python3
"""
nuevas.py — Las dos plantillas nuevas del Kit de Escandallos Pro v2.1.

    12-test-de-rendimiento.xlsx        (Instrucciones · Test de despiece · Test de cocción)
    13-lista-precios-ingredientes.xlsx (Instrucciones · Lista de precios)

Especificación: `scripts/productos-digitales/recipe-costing-kit/SPEC.md` §2.0 (D13, D17,
R2T-06/07/08, R2-02/09/10/13). Datos de los ejemplos, con su fuente, en
`datos_ejemplos.json` (nada de cifras en el código).

Convenciones del kit v2.0 (se IMPORTAN de `kit-escandallos-v2_0/motor.py`, no se copian):
celdas de entrada en verde E8F5E9 y calculadas sin relleno; divisiones con IFERROR;
funciones que pycel evalúa (nada de COUNTA: COUNTIF(rango,"?*")); protección SIN
contraseña con las entradas desbloqueadas (`motor.proteger` vía `motor.cerrar`); A4,
ajuste a una página de ancho y pie «AI Chef Pro · aichef.pro · Página &P de &N»
(`motor.print_setup`); fechas con numFmtId 14 (fecha corta del sistema).

NO guarda caché de fórmulas: `inject_cache.py` va SIEMPRE al final, en main.py.

Uso suelto (depuración):  python3 nuevas.py <carpeta_destino> [<carpeta_origen_fichas>]
"""
import collections
import copy
import datetime
import json
import os
import sys
import unicodedata

AQUI = os.path.dirname(os.path.abspath(__file__))
SCRIPTS = os.path.dirname(AQUI)
ROOT = os.path.dirname(os.path.dirname(SCRIPTS))
sys.path.insert(0, os.path.join(SCRIPTS, 'kit-escandallos-v2_0'))

import openpyxl                                                # noqa: E402
from openpyxl.formatting.rule import CellIsRule                # noqa: E402
from openpyxl.styles import (Alignment, Border, Font,          # noqa: E402
                             PatternFill, Side)
from openpyxl.utils import get_column_letter                   # noqa: E402

import motor                                                   # noqa: E402

ORIGEN = os.path.join(ROOT, 'astro-site', 'public', 'dl', 'kit-escandallos')
DATOS = os.path.join(AQUI, 'datos_ejemplos.json')

F12 = '12-test-de-rendimiento.xlsx'
F13 = '13-lista-precios-ingredientes.xlsx'
SUBJECT = 'Kit de Escandallos Pro · v2.1'
FMT_FECHA = 'mm-dd-yy'          # = numFmtId 14 (fecha corta del sistema), D13
FMT_EUR4 = '#,##0.0000 €'
FMT_KG = '0.000'
VERDE_OK_BG, VERDE_OK_FG = 'C6EFCE', '006100'   # semáforo del 09
GRIS_NOTA = '606060'

MESES = ['enero', 'febrero', 'marzo', 'abril', 'mayo', 'junio', 'julio',
         'agosto', 'septiembre', 'octubre', 'noviembre', 'diciembre']


def version_line(hoy=None):
    hoy = hoy or datetime.date.today()
    return (f'Versión 2.1 · {MESES[hoy.month - 1]} {hoy.year} · '
            'aichef.pro/kit-escandallos · info@aichef.pro')


def cargar_datos():
    with open(DATOS, encoding='utf-8') as fh:
        return json.load(fh)


def es_num(x, dec=2):
    """Número en formato español: 1.234,56."""
    s = f'{x:,.{dec}f}'
    return s.replace(',', '§').replace('.', ',').replace('§', '.')


# ==========================================================================
# Estilo común (paleta y formatos: motor)
# ==========================================================================
FINO = Side(style='thin', color='BFBFBF')
BORDE = Border(left=FINO, right=FINO, top=FINO, bottom=FINO)


def _verde():
    return PatternFill('solid', fgColor=motor.VERDE)


def _relleno(color):
    return PatternFill('solid', fgColor=color)


def titulo_hoja(ws, texto, ultima_col):
    c = ws.cell(row=1, column=1, value=texto)
    c.font = Font(bold=True, size=16, color=motor.ORO)
    c.alignment = Alignment(vertical='center')
    ws.merge_cells(f'A1:{ultima_col}1')
    ws.row_dimensions[1].height = 24


def nota_hoja(ws, fila, texto, ultima_col, ancho_chars, size=11, color='888888'):
    c = ws.cell(row=fila, column=1, value=texto)
    c.font = Font(size=size, color=color, italic=(size < 11))
    c.alignment = Alignment(wrap_text=True, vertical='top')
    ws.merge_cells(f'A{fila}:{ultima_col}{fila}')
    lineas = max(1, -(-len(texto) // max(20, ancho_chars)))
    ws.row_dimensions[fila].height = max(15.0, (13.5 if size >= 11 else 12.0) * lineas + 4)


def banda(ws, fila, texto, ultima_col):
    c = ws.cell(row=fila, column=1, value=texto)
    c.font = Font(bold=True, color='FFFFFF')
    c.fill = _relleno(motor.CAB)
    c.alignment = Alignment(vertical='center')
    ws.merge_cells(f'A{fila}:{ultima_col}{fila}')
    ws.row_dimensions[fila].height = 20


def etiqueta(ws, fila, texto, bold=False, size=11, relleno=None, unir=True):
    c = ws.cell(row=fila, column=1, value=texto)
    c.font = Font(bold=bold, size=size)
    c.alignment = Alignment(vertical='center', wrap_text=True)
    if unir:
        ws.merge_cells(start_row=fila, start_column=1, end_row=fila, end_column=2)
    if relleno and unir:
        for col in (1, 2):
            ws.cell(row=fila, column=col).fill = _relleno(relleno)
    return c


def entrada(ws, coord, valor, fmt=None, align='center'):
    c = ws[coord]
    c.value = valor
    c.fill = _verde()
    c.border = BORDE
    c.font = Font(size=11)
    c.alignment = Alignment(horizontal=align, vertical='center')
    if fmt:
        c.number_format = fmt
    return c


def calculada(ws, coord, formula, fmt, bold=False, size=11, relleno=None):
    c = ws[coord]
    c.value = formula
    c.number_format = fmt
    c.border = BORDE
    c.font = Font(bold=bold, size=size)
    c.alignment = Alignment(horizontal='center', vertical='center')
    if relleno:
        c.fill = _relleno(relleno)
    return c


def nota_celda(ws, fila, texto, desde='D', hasta='F', ancho=40):
    c = ws[f'{desde}{fila}']
    c.value = texto
    c.font = Font(size=9, italic=True, color=GRIS_NOTA)
    c.alignment = Alignment(wrap_text=True, vertical='center')
    ws.merge_cells(f'{desde}{fila}:{hasta}{fila}')
    lineas = max(1, -(-len(texto) // ancho))
    alto = max(18.0, 12.0 * lineas + 6)
    actual = ws.row_dimensions[fila].height or 0
    ws.row_dimensions[fila].height = max(actual, alto)


def metadatos(wb, titulo):
    p = wb.properties
    p.creator = 'AI Chef Pro'
    p.lastModifiedBy = 'AI Chef Pro'
    p.title = f'{titulo} · Kit de Escandallos Pro'
    p.subject = SUBJECT
    p.description = 'aichef.pro/kit-escandallos'
    p.keywords = 'kit escandallos, AI Chef Pro'
    p.category = 'AI Chef Pro · Productos digitales'
    # R1 T7: openpyxl escribe `created` con «Z» (UTC) y pone `modified` en UTC al
    # guardar; con la hora local, la creación salía 2 h DESPUÉS de la modificación.
    ahora = datetime.datetime.now(datetime.timezone.utc).replace(tzinfo=None,
                                                                 microsecond=0)
    p.created = ahora
    p.modified = ahora


# ==========================================================================
# Instrucciones (misma maqueta que las de 01-11: col A de 3, col B de 80)
# ==========================================================================
def hoja_instrucciones(wb, titulo, bloques):
    ws = wb.active
    ws.title = 'Instrucciones'
    ws.sheet_properties.tabColor = motor.ORO
    ws.column_dimensions['A'].width = 3
    ws.column_dimensions['B'].width = 80
    c = ws.cell(row=2, column=2, value=f'📋 {titulo}')
    c.font = Font(name='Calibri', size=18, bold=True, color=motor.ORO)
    c.alignment = Alignment(horizontal='left', vertical='center', wrap_text=True)
    c = ws.cell(row=3, column=2, value='Kit de Escandallos Pro — AI Chef Pro')
    c.font = Font(name='Calibri', size=11, color='888888')
    c = ws.cell(row=4, column=2, value='www.aichef.pro')
    c.font = Font(name='Calibri', size=10, color='888888')

    f_tit = Font(name='Calibri', size=12, bold=True, color='1A1A1A')
    f_lin = Font(name='Calibri', size=11, color='333333')
    al = Alignment(vertical='top', wrap_text=True)
    fila = 6
    for tit, lineas in bloques:
        c = ws.cell(row=fila, column=2, value=tit)
        c.font, c.alignment = f_tit, al
        fila += 2
        for t in lineas:
            c = ws.cell(row=fila, column=2, value=t)
            c.font, c.alignment = copy.copy(f_lin), copy.copy(al)
            fila += 1
        fila += 1
    for t in (motor.PIE, motor.COPY, version_line()):
        c = ws.cell(row=fila, column=2, value=t)
        c.font, c.alignment = copy.copy(f_lin), copy.copy(al)
        fila += 2
    motor.print_setup(ws, None, landscape=False)
    return ws


# ==========================================================================
# 12 · Test de rendimiento
# ==========================================================================
# Coordenadas: una sola fuente para el generador, main.py y verificar.py.
DESP = {
    'producto': 'C5', 'proveedor': 'C6', 'fecha': 'C7', 'ap': 'C8', 'precio': 'C9',
    'coste_ap': 'C10', 'unidades': 'C11',
    'tabla_cab': 13, 'd0': 14, 'd1': 25, 'perdida': 26,
    'peso_util': 'C29', 'valor_sub': 'C30', 'neto': 'C31', 'rend': 'C32',
    'coste_kg_util': 'C33', 'factor': 'C34', 'merma': 'C35', 'porcion': 'C36',
    'coste_porcion': 'C37', 'nota': 39,
}
COC = {
    'elaboracion': 'C5', 'fecha': 'C6', 'crudo': 'C7', 'coste_kg': 'C8',
    'cocinado': 'C9', 'merma_desp': 'C10', 'unidades': 'C11', 'porcion': 'C12',
    'perdida': 'C15', 'coste_kg_coc': 'C16', 'coste_porcion': 'C17',
    'merma_comb': 'C18', 'nota': 20,
}
TIPOS = ['Útil', 'Subproducto', 'Desecho']


def hoja_despiece(wb, d):
    ws = wb.create_sheet('Test de despiece')
    ws.sheet_properties.tabColor = motor.ORO
    motor._cabecera(ws, DESP['tabla_cab'],
                    ['Pieza', 'Tipo', 'Peso (kg)', 'Valor de mercado (€/kg)',
                     'Valor (€)', '% del peso bruto'],
                    [44, 15, 13, 17, 13, 12])
    for col in range(1, 7):
        c = ws.cell(row=DESP['tabla_cab'], column=col)
        c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    ws.row_dimensions[DESP['tabla_cab']].height = 30

    titulo_hoja(ws, d['titulo'], 'F')
    nota_hoja(ws, 2, 'Pesa la pieza tal como llega y cada parte que sale al '
              'limpiarla. Las celdas verdes son editables; el resto son fórmulas.',
              'F', 110)

    banda(ws, 4, 'DATOS DE LA COMPRA', 'F')
    etiqueta(ws, 5, 'Producto', bold=True)
    entrada(ws, 'C5', d['producto'], align='left')
    ws.merge_cells('C5:F5')
    etiqueta(ws, 6, 'Proveedor', bold=True)
    entrada(ws, 'C6', None, align='left')
    ws.merge_cells('C6:F6')
    etiqueta(ws, 7, 'Fecha del test', bold=True)
    entrada(ws, 'C7', None, FMT_FECHA)
    # R1 CH-09: «peso bruto», no «AP» (en España nadie dice AP; el blog propio
    # tampoco). El EN sí usa AP (SPEC D8).
    etiqueta(ws, 8, 'Peso bruto de compra (kg)', bold=True)
    entrada(ws, 'C8', d['peso_ap_kg'], FMT_KG)
    nota_celda(ws, 8, 'La pieza entera, tal como llega.')
    etiqueta(ws, 9, 'Precio por kg de compra (€/kg)', bold=True)
    entrada(ws, 'C9', d['precio_ap_eur_kg'], motor.FMT_EUR)
    nota_celda(ws, 9, 'El de la factura: el mismo que pones en la ficha.')
    etiqueta(ws, 10, 'Coste bruto de compra (€)', bold=True)
    calculada(ws, 'C10', '=IFERROR(C8*C9,"")', motor.FMT_EUR, bold=True)
    nota_celda(ws, 10, 'Peso × precio.')
    etiqueta(ws, 11, 'Gramos por kilo (para la porción)', bold=True)
    entrada(ws, 'C11', d['unidades_porcion_por_unidad_peso'], '#,##0')
    nota_celda(ws, 11, 'Déjalo en 1.000 si pesas la pieza en kg y la porción en g.')

    # ---- tabla de componentes (12 filas) ----------------------------------
    d0, d1 = DESP['d0'], DESP['d1']
    comps = d['componentes']
    for i, r in enumerate(range(d0, d1 + 1)):
        comp = comps[i] if i < len(comps) else None
        entrada(ws, f'A{r}', comp['pieza'] if comp else None, align='left')
        entrada(ws, f'B{r}', comp['tipo'] if comp else None)
        entrada(ws, f'C{r}', comp['peso_kg'] if comp else None, FMT_KG)
        entrada(ws, f'D{r}', comp['valor_eur_kg'] if comp else None, motor.FMT_EUR)
        calculada(ws, f'E{r}',
                  f'=IF(B{r}="Subproducto",IFERROR(C{r}*D{r},""),"")', motor.FMT_EUR)
        calculada(ws, f'F{r}', f'=IF(C{r}="","",IFERROR(C{r}/$C$8,""))',
                  motor.FMT_PCT1)
    dv = motor._dv_lista('"' + ','.join(TIPOS) + '"',
                         'Elige Útil, Subproducto o Desecho.')
    ws.add_data_validation(dv)
    dv.add(f'B{d0}:B{d1}')

    rp = DESP['perdida']
    c = ws.cell(row=rp, column=1, value='Pérdida de corte (goteo y recortes que se pierden)')
    c.font = Font(italic=True)
    c.border = BORDE
    c = ws.cell(row=rp, column=2, value='Desecho')
    c.alignment = Alignment(horizontal='center')
    c.border = BORDE
    # ROUND: sin él, 2,8 − (2,1 + 0,16 + 0,54) da −4e-16 en coma flotante y
    # pintaría de rojo una pérdida que es cero. El «+0» convierte el −0 que deja
    # ROUND en 0 (si no, la caché guarda «-0.0» y un visor puede enseñar «-0,000»).
    calculada(ws, f'C{rp}', f'=IFERROR(ROUND($C$8-SUM(C{d0}:C{d1}),6)+0,"")', FMT_KG)
    ws.cell(row=rp, column=4).border = BORDE
    ws.cell(row=rp, column=5).border = BORDE
    calculada(ws, f'F{rp}', f'=IF(C{rp}="","",IFERROR(C{rp}/$C$8,""))', motor.FMT_PCT1)
    # Negativa = algo se ha pesado dos veces (lo dicen las Instrucciones).
    ws.conditional_formatting.add(
        f'C{rp}', CellIsRule(operator='lessThan', formula=['0'],
                             fill=_relleno(motor.ROJO_BG),
                             font=Font(color=motor.ROJO_FG)))

    # ---- resultado ------------------------------------------------------
    banda(ws, 28, 'RESULTADO', 'F')
    filas = [
        (29, 'Peso útil (kg)', f'=SUMIF(B{d0}:B{d1},"Útil",C{d0}:C{d1})', FMT_KG,
         False, None, 'Suma de las partes «Útil».'),
        (30, 'Valor de los subproductos (€)', f'=SUM(E{d0}:E{d1})', motor.FMT_EUR,
         False, None, 'Lo que valen las partes que aprovechas.'),
        (31, 'Coste neto útil (€)', '=IFERROR(C10-C30,"")', motor.FMT_EUR,
         False, None, 'Coste bruto de compra − valor de los subproductos.'),
        (32, 'Rendimiento bruto (%) · informativo', '=IFERROR(C29/C8,"")',
         motor.FMT_PCT1, False, None, 'Peso útil ÷ peso bruto. No va a la ficha.'),
        (33, 'COSTE POR KG ÚTIL (€/kg)', '=IFERROR(C31/C29,"")', motor.FMT_EUR,
         True, motor.CREMA, 'Lo que te cuesta cada kilo que sirves.'),
        (34, 'Factor de coste', '=IFERROR(C33/C9,"")', '0.000',
         False, None, 'Precio nuevo × factor = nuevo coste por kg útil.'),
        (35, 'MERMA % PARA TU ESCANDALLO', '=IFERROR(1-C9/C33,"")', motor.FMT_PCT1,
         True, motor.ORO, 'La ÚNICA cifra que va a la ficha, a «Merma (%)».'),
    ]
    for fila, rot, f, fmt, bold, rell, nota in filas:
        etiqueta(ws, fila, rot, bold=bold, size=(12 if rell == motor.ORO else 11),
                 relleno=rell)
        calculada(ws, f'C{fila}', f, fmt, bold=bold,
                  size=(12 if rell == motor.ORO else 11), relleno=rell)
        nota_celda(ws, fila, nota)
    etiqueta(ws, 36, 'Porción estándar (g)', bold=True)
    entrada(ws, 'C36', d['porcion_g'], '#,##0')
    nota_celda(ws, 36, 'En crudo y ya limpia: la ficha 01 usa 200 g.')
    etiqueta(ws, 37, 'COSTE POR PORCIÓN (€)', bold=True, size=12, relleno=motor.ORO)
    calculada(ws, 'C37', '=IFERROR(C36/C11*C33,"")', motor.FMT_EUR, bold=True,
              size=12, relleno=motor.ORO)
    nota_celda(ws, 37, 'Porción × coste por kg útil. La ficha dará lo mismo.')

    nota_hoja(ws, DESP['nota'], d['nota_hoja'], 'F', 125, size=9, color=GRIS_NOTA)
    return ws


def hoja_coccion(wb, d):
    ws = wb.create_sheet('Test de cocción')
    ws.sheet_properties.tabColor = motor.ORO
    for col, ancho in zip('ABCDEF', [44, 15, 13, 17, 13, 12]):
        ws.column_dimensions[col].width = ancho
    titulo_hoja(ws, d['titulo'], 'F')
    nota_hoja(ws, 2, 'Pesa en crudo lo que vas a cocinar, cocínalo como en servicio y '
              'pésalo antes de racionar. Las celdas verdes son editables.', 'F', 110)

    banda(ws, 4, 'DATOS DEL TEST', 'F')
    etiqueta(ws, 5, 'Elaboración', bold=True)
    entrada(ws, 'C5', d['elaboracion'], align='left')
    ws.merge_cells('C5:F5')
    etiqueta(ws, 6, 'Fecha del test', bold=True)
    entrada(ws, 'C6', None, FMT_FECHA)
    etiqueta(ws, 7, 'Peso crudo útil (kg)', bold=True)
    entrada(ws, 'C7', d['peso_crudo_kg'], FMT_KG)
    nota_celda(ws, 7, 'Ya limpio: lo que entra en la plancha.')
    # R1 CH-04: el precio de entrada es el de la FACTURA (el de la ficha) y la
    # merma de despiece se aplica en C16. Así el coste por porción cocinada
    # (C17) coincide SIEMPRE con la ficha que lleve la merma combinada (C18).
    etiqueta(ws, 8, 'Precio de compra por kg (€/kg)', bold=True)
    entrada(ws, 'C8', d['coste_crudo_eur_kg'], motor.FMT_EUR)
    nota_celda(ws, 8, 'El de la factura: el mismo que pones en la ficha.')
    etiqueta(ws, 9, 'Peso cocinado (kg)', bold=True)
    entrada(ws, 'C9', d['peso_cocinado_kg'], FMT_KG)
    nota_celda(ws, 9, 'Recién hecho, antes de racionar.')
    etiqueta(ws, 10, 'Merma de despiece de esta pieza (%)', bold=True)
    entrada(ws, 'C10', d['merma_despiece'], motor.FMT_PCT1)
    nota_celda(ws, 10, 'La MERMA % del test de despiece; 0 % si la compras lista, '
                       'como esta picada.')
    etiqueta(ws, 11, 'Gramos por kilo (para la porción)', bold=True)
    entrada(ws, 'C11', d['unidades_porcion_por_unidad_peso'], '#,##0')
    nota_celda(ws, 11, 'Déjalo en 1.000 si pesas en kg y racionas en g.')
    etiqueta(ws, 12, 'Porción cocinada (g)', bold=True)
    entrada(ws, 'C12', d['porcion_cocinada_g'], '#,##0')
    nota_celda(ws, 12, 'La que sirves: 180 g crudos de la ficha 08 × 0,74.')

    banda(ws, 14, 'RESULTADO', 'F')
    filas = [
        (15, 'PÉRDIDA POR COCCIÓN (%)', '=IFERROR(1-C9/C7,"")', motor.FMT_PCT1,
         True, motor.CREMA, '1 − peso cocinado ÷ peso crudo.'),
        (16, 'Coste por kg cocinado (€/kg)', '=IFERROR(C8/((1-C10)*(1-C15)),"")',
         motor.FMT_EUR, False, None,
         'Precio ÷ ((1 − merma de despiece) × (1 − pérdida por cocción)).'),
        (17, 'COSTE POR PORCIÓN COCINADA (€)', '=IFERROR(C12/C11*C16,"")',
         motor.FMT_EUR, True, motor.ORO, 'Porción × coste por kg cocinado.'),
        (18, 'MERMA COMBINADA PARA LA FICHA (%)', '=IFERROR(1-(1-C10)*(1-C15),"")',
         motor.FMT_PCT1, True, motor.ORO,
         'Si en la ficha escribes el peso COCINADO.'),
    ]
    for fila, rot, f, fmt, bold, rell, nota in filas:
        etiqueta(ws, fila, rot, bold=bold, size=(12 if rell == motor.ORO else 11),
                 relleno=rell)
        calculada(ws, f'C{fila}', f, fmt, bold=bold,
                  size=(12 if rell == motor.ORO else 11), relleno=rell)
        nota_celda(ws, fila, nota)

    nota_hoja(ws, COC['nota'], d['nota_hoja'], 'F', 125, size=9, color=GRIS_NOTA)
    return ws


def merma_python(d):
    """La misma cuenta que la hoja, en Python: sirve para el texto de las
    Instrucciones y como contraprueba del valor que calcula pycel."""
    ap = d['peso_ap_kg'] * d['precio_ap_eur_kg']
    util = sum(c['peso_kg'] for c in d['componentes'] if c['tipo'] == 'Útil')
    sub = sum(c['peso_kg'] * c['valor_eur_kg'] for c in d['componentes']
              if c['tipo'] == 'Subproducto')
    coste_kg_util = (ap - sub) / util
    return 1 - d['precio_ap_eur_kg'] / coste_kg_util


def raciones_python(d):
    """Raciones que salen de la pieza (despiece) y de la tanda (cocción), para la
    línea de Instrucciones de R1 CH-12 (la fila calculada queda pendiente de John)."""
    util = sum(c['peso_kg'] for c in d['despiece']['componentes'] if c['tipo'] == 'Útil')
    r_desp = util * d['despiece']['unidades_porcion_por_unidad_peso'] / d['despiece']['porcion_g']
    c = d['coccion']
    r_coc = c['peso_cocinado_kg'] * c['unidades_porcion_por_unidad_peso'] / c['porcion_cocinada_g']
    return util, r_desp, r_coc


def pct(x, dec=1):
    """0,235437 → «23,5 %» (sin decimales si es entero a esa precisión)."""
    v = round(x * 100, dec)
    if dec and abs(v - round(v)) < 1e-9:
        return f'{round(v):.0f} %'
    return es_num(v, dec) + ' %'


def instrucciones_12(datos):
    d = datos['despiece']
    merma = merma_python(d)
    generica = dict((m[0], m[1]) for m in motor.MERMAS)['Carne roja']
    util, r_desp, r_coc = raciones_python(datos)
    limpieza = 1 - util / d['peso_ap_kg']
    c = datos['coccion']
    return [
        ('Para Qué Sirve', [
            '▸ La merma de la hoja «Mermas» de cada ficha es un punto de partida. La tuya depende del proveedor, de la pieza y de cómo la limpias, y la única forma de conocerla es pesar. Esta plantilla convierte ese pesaje en la cifra que necesita tu escandallo.',
            '▸ Pestaña «Test de despiece»: una pieza que compras entera (un solomillo, un lomo, un pescado) y lo que sale de ella al limpiarla.',
            '▸ Pestaña «Test de cocción»: lo que pierde una elaboración al cocinarla, para cuando racionas en peso cocinado.',
            '▸ Las celdas VERDES son las editables; el resto son fórmulas.',
        ]),
        ('Test de Despiece, Paso a Paso', [
            '▸ 1. Pesa la pieza tal como llega (peso bruto) y escribe el precio por kg de la factura.',
            '▸ 2. Límpiala y pesa por separado cada parte. Márcala como «Útil» (lo que vas a servir), «Subproducto» (lo que aprovechas en otra elaboración) o «Desecho» (lo que tiras).',
            '▸ 3. A cada subproducto ponle su valor de mercado por kg: lo que te costaría comprar esa carne para el uso que le das (el cordón para una picada o un fondo, las puntas para un salteado). Ese valor se descuenta del coste de la pieza. Si no lo vas a aprovechar, márcalo como «Desecho».',
            '▸ 4. La fila «Pérdida de corte» calcula sola lo que no has pesado: el goteo y los recortes que se quedan en la tabla. Si sale negativa (en rojo), has pesado algo dos veces.',
            '▸ 5. Lleva a la ficha la MERMA % PARA TU ESCANDALLO: escríbela en la columna «Merma (%)» de ese ingrediente y deja el precio por kg de compra de siempre. La ficha dará el mismo coste por ración que este test.',
        ]),
        ('Cómo Leer el Resultado', [
            '▸ Peso útil: la suma de las partes marcadas «Útil».',
            '▸ Coste neto útil: lo que pagaste por la pieza menos lo que valen los subproductos.',
            '▸ Rendimiento bruto (%): peso útil ÷ peso bruto. Es informativo: no tiene en cuenta el valor de los subproductos, así que no lo lleves a la ficha.',
            '▸ COSTE POR KG ÚTIL: lo que te cuesta de verdad cada kilo que sirves.',
            '▸ Factor de coste: coste por kg útil ÷ precio por kg de compra. Si el proveedor sube el precio de la pieza, multiplica el precio nuevo por este factor y tendrás el nuevo coste por kg útil sin repetir el test, mientras no cambien la pieza ni la forma de limpiarla.',
            '▸ MERMA % PARA TU ESCANDALLO: 1 − precio de compra ÷ coste por kg útil. Ya descuenta lo que valen los subproductos y es la única cifra que va a la ficha.',
            '▸ COSTE POR PORCIÓN: porción estándar × coste por kg útil. La celda verde «Gramos por kilo» vale 1.000 porque la pieza se pesa en kg y la porción en g; si pesas y racionas en la misma unidad, ponla a 1.',
            '▸ ¿Cuántas raciones salen de la pieza? Peso útil × 1.000 ÷ porción estándar. En el ejemplo, '
            + es_num(util, 3) + ' × 1.000 ÷ ' + es_num(d['porcion_g'], 0) + ' = '
            + es_num(r_desp, 1) + ' raciones: es la cuenta para saber cuántas piezas pedir.',
        ]),
        ('Test de Cocción', [
            '▸ Pesa en crudo lo que vas a cocinar, ya limpio, y escribe el precio por kg de la factura, el mismo de la ficha. Si viene de una pieza que despiezas, escribe también su merma de despiece (la MERMA % PARA TU ESCANDALLO del test de despiece); si la compras lista, déjala en 0 %.',
            '▸ Cocínalo como en servicio y pésalo recién hecho, antes de racionar. PÉRDIDA POR COCCIÓN = 1 − peso cocinado ÷ peso crudo.',
            '▸ Si en la ficha escribes la cantidad en peso CRUDO, usa como merma la del despiece, sin más. Si la escribes en peso COCINADO (la porción que sale al plato), usa la MERMA COMBINADA PARA LA FICHA: 1 − (1 − merma de despiece) × (1 − pérdida por cocción). Con cualquiera de las dos, la ficha da el mismo coste por porción que este test.',
            '▸ Raciones de la tanda: peso cocinado × 1.000 ÷ porción cocinada. En el ejemplo, '
            + es_num(c['peso_cocinado_kg'], 3) + ' × 1.000 ÷ ' + es_num(c['porcion_cocinada_g'], 0)
            + ' = ' + es_num(r_coc, 0) + ' hamburguesas.',
            '▸ Duplica la pestaña para otro test: clic derecho sobre la pestaña → Mover o copiar → Crear una copia.',
        ]),
        ('Los Ejemplos Cargados', [
            '▸ Despiece: el solomillo de ternera entero, con cordón, que compra la plantilla 01, a '
            + es_num(d['precio_ap_eur_kg']) + ' €/kg. Es un ejemplo orientativo, no un test pesado: una pieza de '
            + es_num(d['peso_ap_kg'], 1) + ' kg (las enteras pesan entre 2,5 y 3,2 kg) con una merma de limpieza del '
            + pct(limpieza, 0) + ', la referencia habitual para el solomillo entero con cordón. El reparto entre el cordón que se aprovecha y lo que se tira, y el valor del cordón, son estimaciones: pon los tuyos.',
            f'▸ Con ese despiece, la merma del solomillo para la ficha es del {pct(merma)} (el {pct(limpieza, 0)} de limpieza menos lo que vale el cordón) y no del {pct(generica, 0)} genérico de «Carne roja»: es la que lleva ahora la plantilla 01 en la fila del solomillo.',
            '▸ Cocción: la smash burger de la plantilla 08, 10 hamburguesas de 180 g de mezcla 80/20. El peso cocinado aplica el rendimiento de referencia del USDA para picada de vacuno con un 20 % de grasa (0,74 kg cocinados por kg crudo). No es una medida: pesa la tuya, porque la plancha, el grosor y el punto lo cambian.',
            '▸ La picada se compra lista, así que su merma de despiece es 0 %, y la ficha 08 lleva también un 0 % en esa fila: con los 180 g en crudo cuesta lo mismo en la ficha que aquí.',
            '▸ Son datos de ejemplo: escribe encima de las celdas verdes o vacíalas con Supr.',
        ]),
        ('Protección de la Hoja', [
            '▸ Las hojas están protegidas SIN contraseña para que no borres una fórmula sin querer: se escriben únicamente las celdas verdes. Revisar → Desproteger hoja para tocar el resto; no pide contraseña.',
        ]),
    ]


def _cerrar_comun(wb, fname, informe):
    antes = motor.VERSION_LINE
    motor.VERSION_LINE = version_line()
    try:
        motor.cerrar(wb, fname, informe)
    finally:
        motor.VERSION_LINE = antes


def generar_12(destino, datos, informe=None):
    informe = informe if informe is not None else []
    wb = openpyxl.Workbook()
    hoja_instrucciones(wb, 'Test de Rendimiento — Despiece y Cocción',
                       instrucciones_12(datos))
    hoja_despiece(wb, datos['despiece'])
    hoja_coccion(wb, datos['coccion'])
    metadatos(wb, 'Test de Rendimiento: Despiece y Cocción')
    _cerrar_comun(wb, F12, informe)
    # motor.cerrar deja el área de impresión hasta la columna J como mínimo;
    # aquí la tabla acaba en F y las 4 columnas vacías encogerían el papel.
    # R1 CH-07: motor.cerrar las pone en horizontal (≥ 6 columnas) y el despiece
    # se partía en dos A4 con la banda RESULTADO sola al pie de la primera. En
    # vertical, con ajuste a 1 página de ancho, cada test cabe en una hoja.
    for ws in wb.worksheets[1:]:
        ws.print_area = f'A1:F{ws.max_row}'
        ws.page_setup.orientation = 'portrait'
    path = os.path.join(destino, F12)
    wb.save(path)
    return path


# ==========================================================================
# 13 · Lista de precios de ingredientes
# ==========================================================================
LISTA = {'resumen_n': 'B4', 'umbral': 'B5', 'alertas': 'B6', 'mayor': 'B7',
         'cab': 9, 'd0': 10}
COLS13 = ['Ingrediente', 'Categoría', 'Proveedor', 'Formato de compra',
          'Contenido del formato', 'Unidad base', 'Precio del formato (€)',
          'Precio por unidad base (€)', 'Precio anterior por unidad base (€)',
          'Variación (%)', 'Estado', 'Fecha última factura', 'Notas']
ANCHOS13 = [34, 18, 16, 26, 11, 9, 12, 13, 13, 11, 10, 12, 52]
COL_IMPRESION13 = 'L'       # R1 CH-08: Notas (M) no se imprime
CHARS_NOTA13 = 58           # caracteres por línea en M (ancho 52, letra de 9 pt)


def _clave_orden(nombre):
    s = unicodedata.normalize('NFKD', nombre)
    return ''.join(ch for ch in s if not unicodedata.combining(ch)).lower()


def conversiones(origen):
    wb = openpyxl.load_workbook(os.path.join(origen, motor.FICHEROS[0]), data_only=True)
    ws = wb['Conversiones']
    fuera = {}
    r = 5
    while ws.cell(row=r, column=1).value:
        fuera[ws.cell(row=r, column=1).value] = ws.cell(row=r, column=2).value
        r += 1
    return fuera


def corregir_item(nombre, categoria, datos=None):
    """R1 CH-10: correcciones de TEXTO de las filas de ejemplo de 01-08 (nombre y
    categoría). Las mismas que aplica parche_v2_1.py a las fichas; aquí se aplican
    a la LECTURA de las fichas publicadas, para que el 13 ya salga corregido."""
    cf = (datos or cargar_datos()).get('correcciones_fichas', {})
    nombre = cf.get('renombrar', {}).get(nombre, nombre)
    categoria = cf.get('categoria', {}).get(nombre, categoria)
    return nombre, categoria


def leer_fichas(origen, datos=None):
    """Todas las filas con ingrediente de las 8 fichas publicadas (01-08), con las
    correcciones de texto de la v2.1 ya aplicadas (`corregir_item`)."""
    datos = datos or cargar_datos()
    for cat in datos.get('correcciones_fichas', {}).get('categoria', {}).values():
        if cat not in motor.CATEGORIAS:
            raise SystemExit(f'correcciones_fichas: «{cat}» no es una categoría del kit')
    fuera = []
    for fname in motor.FICHEROS:
        path = os.path.join(origen, fname)
        wb = openpyxl.load_workbook(path)
        wbv = openpyxl.load_workbook(path, data_only=True)
        for ws, lay in motor.hojas_escandallo(wb):
            wsv = wbv[ws.title]
            for r in range(lay.d0, lay.d1 + 1):
                nombre = ws.cell(row=r, column=1).value
                if not (isinstance(nombre, str) and nombre.strip()):
                    continue
                precio = wsv.cell(row=r, column=4).value
                if not isinstance(precio, (int, float)):
                    raise SystemExit(f'{fname}:{ws.title}!D{r}: precio sin valor '
                                     f'({precio!r}); ¿falta la caché?')
                nom, cat = corregir_item(nombre.strip(), ws.cell(row=r, column=2).value,
                                         datos)
                fuera.append({
                    'fichero': fname, 'plantilla': fname[:2], 'hoja': ws.title,
                    'fila': r, 'nombre': nom, 'categoria': cat,
                    'ud': ws.cell(row=r, column=3).value, 'precio': float(precio),
                })
    return fuera


def se_usa_en(items):
    """R1 CH-06: «Se usa en: 03 «Postre» · 07 «Tostada Aguacate», «Eggs Benedict»».
    Las pestañas van entre comillas angulares, como en el resto del kit (hay
    nombres con paréntesis: «Cocktail (por persona)»)."""
    por = collections.OrderedDict()
    for it in items:
        hojas = por.setdefault(it['plantilla'], [])
        if it['hoja'] not in hojas:
            hojas.append(it['hoja'])
    return ('Se usa en: '
            + ' · '.join(f'{p} ' + ', '.join(f'«{h}»' for h in hs)
                         for p, hs in sorted(por.items())) + '.')


def filas_lista(origen, datos):
    lp = datos['lista_precios']
    alias, formatos, defecto = lp['alias'], lp['formatos'], lp['formato_por_defecto']
    excluir = lp.get('excluir', {})
    conv = conversiones(origen)
    fichas = leer_fichas(origen, datos)

    nombres_ficha = {f['nombre'] for f in fichas}
    avisos = [f'alias sin uso: {k}' for k in alias if k not in nombres_ficha]
    avisos += [f'exclusión sin uso: {k}' for k in excluir if k not in nombres_ficha]
    excluidas = [f for f in fichas if f['nombre'] in excluir]
    fichas = [f for f in fichas if f['nombre'] not in excluir]

    grupos = collections.OrderedDict()
    for it in fichas:
        grupos.setdefault(alias.get(it['nombre'], it['nombre']), []).append(it)
    avisos += [f'formato sin ingrediente: {k}' for k in formatos if k not in grupos]

    filas, conflictos, mapa = [], [], []
    for canon, items in grupos.items():
        # El precio de la lista es el de la fila cuyo nombre ES el canónico (si
        # existe): «Sal» y no «Sal (para borde)», aunque el 04 salga antes.
        base = next((it for it in items if it['nombre'] == canon), items[0])
        otros = []
        for it in items:
            if it is base:
                continue
            igual = (it['ud'] == base['ud'] and abs(it['precio'] - base['precio']) < 1e-9)
            if igual:
                continue
            # R1 T4/CH-10: un alias puede unir precios distintos del MISMO producto
            # (otro proveedor) con la misma unidad; se avisa en Notas. Con otra
            # unidad es casi seguro un alias equivocado: se aborta.
            if it['ud'] != base['ud']:
                raise SystemExit(
                    f'«{it["nombre"]}» → «{canon}» une unidades distintas '
                    f'({base["precio"]} {base["ud"]} vs {it["precio"]} {it["ud"]})')
            otros.append(it)
        formato, contenido, ub = formatos.get(canon) or defecto[base['ud']]
        clave = f'{base["ud"]}→{ub}'
        if clave not in conv:
            raise SystemExit(f'{canon}: no hay conversión {clave} en «Conversiones»')
        factor = float(conv[clave])
        pb = base['precio'] / factor
        pf = round(pb * contenido, 2)
        if abs(round(pf / contenido, 4) - round(pb, 4)) > 1e-9:
            raise SystemExit(f'{canon}: {pf} € / {contenido} ≠ {pb} (elige otro formato)')
        notas = [se_usa_en(items)]
        anterior = lp['precio_anterior_demo'].get(canon)
        if anterior is not None:
            notas.append(lp['nota_demo'])
        vistos = set()
        for o in otros:
            k = (o['plantilla'], o['precio'], o['ud'])
            if k in vistos:
                continue
            vistos.add(k)
            notas.append(f'La plantilla {o["plantilla"]} lo usa a '
                         f'{es_num(o["precio"])} €/{o["ud"]} (otro proveedor).')
            conflictos.append({'ingrediente': canon, 'nombre_ficha': o['nombre'],
                               'lista': base['precio'], 'ud': base['ud'],
                               'plantilla': o['plantilla'], 'hoja': o['hoja'],
                               'fila': o['fila'], 'ficha': o['precio']})
        filas.append({
            'ingrediente': canon, 'categoria': base['categoria'], 'formato': formato,
            'contenido': contenido, 'ub': ub, 'precio_formato': pf,
            'precio_base': pb, 'anterior': anterior, 'notas': ' '.join(notas),
            'ud_ficha': base['ud'], 'factor': factor,
        })
        for it in items:
            mapa.append({'fichero': it['fichero'], 'hoja': it['hoja'], 'fila': it['fila'],
                         'nombre': it['nombre'], 'canon': canon, 'ud': it['ud'],
                         'precio': it['precio']})
    orden = {c: i for i, c in enumerate(motor.CATEGORIAS)}
    filas.sort(key=lambda f: (orden.get(f['categoria'], 99), _clave_orden(f['ingrediente'])))
    for f in excluidas:
        avisos.append(f'excluido del 13: {f["fichero"][:2]} {f["hoja"]}!A{f["fila"]} '
                      f'{f["nombre"]} ({excluir[f["nombre"]]})')
    return filas, conflictos, mapa, avisos


def instrucciones_13(n, vacias):
    return [
        ('Para Qué Sirve', [
            '▸ Un único sitio con el precio actual de cada ingrediente, pasado a la unidad que usan tus fichas (kg, L o ud). Cuando llega una factura, actualizas aquí el precio y ves de un vistazo qué ha subido más de la cuenta.',
            f'▸ Viene cargada con los {n} ingredientes de las recetas de ejemplo del kit (plantillas 01 a 08) y sus precios, más {vacias} filas vacías al final para lo que te falte. Escribe encima de los ejemplos o vacíalos seleccionando solo las columnas verdes y pulsando Supr: las blancas son fórmulas y la hoja no deja borrarlas.',
            '▸ Las celdas VERDES son las editables; las blancas son fórmulas.',
        ]),
        ('Qué Es Cada Columna', [
            '▸ Ingrediente, Categoría y Proveedor: texto libre. La categoría es la misma que usan las fichas.',
            '▸ Formato de compra: tal como viene en la factura («Saco 25 kg», «Garrafa 5 L», «Estuche 30 huevos»).',
            '▸ Contenido del formato: cuántas unidades base trae (25 en un saco de 25 kg; 0,7 en una botella de 70 cl si la unidad base es el litro).',
            '▸ Unidad base: la unidad en la que quieres el precio (desplegable). Usa kg, L o ud siempre que puedas: son las que entienden todas las fichas.',
            '▸ Precio del formato (€): lo que pagas por él, sin IVA, como en las fichas.',
            '▸ Precio por unidad base (€): precio del formato ÷ contenido, con 4 decimales. Es el número que se lleva a las fichas.',
            '▸ Precio anterior por unidad base (€): el que tienen ahora tus fichas. Actualízalo (copia el «Precio por unidad base» y pégalo aquí con Pegado especial → Valores) solo cuando lleves el precio nuevo a las fichas.',
            '▸ Variación (%): cuánto se ha movido el precio actual respecto al que tienen tus fichas.',
            '▸ Estado: ALERTA (en rojo) si la subida supera el umbral: tus fichas están desfasadas más de la cuenta y toca volver a escandallar. OK si no. Las bajadas nunca dan alerta.',
            '▸ Fecha última factura: para saber de cuándo es cada precio.',
            '▸ Notas: en los ejemplos dice en qué fichas del kit se usa cada ingrediente. Apunta lo mismo con tus fichas: es la lista de sitios a los que llevar un precio nuevo.',
        ]),
        ('Umbral de Alerta y Resumen', [
            '▸ La celda verde «Umbral de alerta (%)» viene al 5 %: si el precio actual supera en más de eso al que tienen tus fichas, el ingrediente pasa a ALERTA. Una subida exactamente igual al umbral no la dispara.',
            '▸ Así, varias subidas pequeñas seguidas del mismo proveedor acaban saltando: la alerta compara siempre con el precio de tus fichas, no con la última factura.',
            '▸ Arriba tienes el resumen: nº de ingredientes, nº de alertas y la mayor subida de la lista.',
            '▸ En el ejemplo, dos filas llevan un precio anterior marcado «Ilustrativo» en Notas, para que veas una ALERTA y un OK. Bórralos cuando empieces con tus precios.',
        ]),
        ('Cómo Llevar un Precio a la Ficha', [
            '▸ Las fichas NO leen esta lista por su cuenta: cada plantilla es un libro independiente, y enlazar libros entre sí se rompe en cuanto mueves un fichero o lo abres en Google Sheets.',
            '▸ 1. Copia la celda «Precio por unidad base» del ingrediente.',
            '▸ 2. En la ficha, sobre «Precio/Ud (€)», pega con Pegado especial → Valores: en Excel, Ctrl+Alt+V y luego V (o Inicio → Pegar → Pegar valores); en Google Sheets, Ctrl+Mayús+V. Un pegado normal arrastraría la fórmula: el precio y el coste de esa línea se quedarían en blanco sin avisar, y el coste del plato saldría más bajo. Comprueba que «Precio/Ud (€)» muestra el número.',
            '▸ 3. En «Ud. Compra» de la ficha pon la unidad base de la lista. Ejemplo: los huevos camperos se compran por docena en la plantilla 03 (3,60 €/docena); en esta lista van por unidad (0,3000 €/ud). Si pegas 0,30, cambia «Ud. Compra» a «ud».',
            '▸ 4. Copia también el precio nuevo a «Precio anterior» (Pegado especial → Valores): desde ese momento es el que tienen tus fichas.',
            '▸ En la plantilla 04 (cócteles), el precio de los destilados, el vino y el espumoso viene de su pestaña «Formatos de Compra»: allí basta con actualizar el precio de la botella.',
        ]),
        ('Filtrar y Ordenar', [
            '▸ La tabla lleva autofiltro y funciona con la hoja protegida: usa las flechas de la cabecera para filtrar por categoría, proveedor o estado. Para ver lo que más ha subido, filtra «Estado» por ALERTA.',
            '▸ Para ordenar, primero Revisar → Desproteger hoja (no pide contraseña) y ordena siempre desde las flechas de la cabecera, nunca una columna suelta: así cada precio sigue en la fila de su ingrediente. Si ordenas «Variación (%)» de mayor a menor, quita antes (Vacías) en el filtro de esa columna, o las filas sin precio anterior se quedan arriba. Al terminar, Revisar → Proteger hoja.',
        ]),
        ('Protección de la Hoja', [
            '▸ La hoja está protegida SIN contraseña para que no borres una fórmula sin querer: se escriben las celdas verdes, y las blancas («Precio por unidad base», «Variación» y «Estado») están bloqueadas. Revisar → Desproteger hoja para tocar el resto; no pide contraseña.',
            '▸ ¿Te quedas sin filas? Revisar → Desproteger hoja, inserta filas DENTRO de la tabla (encima de la última) y copia en ellas las columnas blancas de la fila de arriba. Así el resumen, el filtro y la impresión las recogen.',
            '▸ Si con la hoja desprotegida borras una fórmula, cópiala de cualquiera de las filas vacías del final.',
        ]),
    ]


def hoja_lista(wb, filas, datos):
    lp = datos['lista_precios']
    ws = wb.create_sheet('Lista de precios')
    ws.sheet_properties.tabColor = motor.ORO
    cab = LISTA['cab']
    motor._cabecera(ws, cab, COLS13, ANCHOS13)
    for col in range(1, len(COLS13) + 1):
        ws.cell(row=cab, column=col).alignment = Alignment(
            horizontal='center', vertical='center', wrap_text=True)
    ws.row_dimensions[cab].height = 44
    ult = get_column_letter(len(COLS13))
    # R1 CH-08: se imprime hasta «Fecha última factura» (L), sin Notas; el título y
    # la nota se combinan hasta L para que no se corten en el papel.
    titulo_hoja(ws, 'Lista de Precios de Ingredientes', COL_IMPRESION13)
    nota_hoja(ws, 2, 'Un precio por ingrediente, pasado a la unidad que usan tus fichas. '
              'Las celdas verdes son editables. Para llevar un precio a una ficha, copia '
              '«Precio por unidad base» y pégalo con Pegado especial → Valores (Excel: '
              'Ctrl+Alt+V y luego V).', COL_IMPRESION13, 180)

    d0 = LISTA['d0']
    d1 = d0 + len(filas) + lp['filas_vacias'] - 1

    # ---- resumen -------------------------------------------------------
    etiqueta(ws, 4, 'Nº de ingredientes', bold=True, unir=False)
    calculada(ws, 'B4', f'=COUNTIF(A{d0}:A{d1},"?*")', '0', bold=True)
    etiqueta(ws, 5, 'Umbral de alerta (%)', bold=True, unir=False)
    entrada(ws, 'B5', lp['umbral_alerta'], motor.FMT_PCT1)
    nota_celda(ws, 5, 'Una subida mayor pone el ingrediente en ALERTA.', 'C', 'F', 60)
    etiqueta(ws, 6, 'Nº de alertas', bold=True, unir=False)
    calculada(ws, 'B6', f'=COUNTIF(K{d0}:K{d1},"ALERTA")', '0', bold=True)
    etiqueta(ws, 7, 'Mayor subida (%)', bold=True, unir=False)
    calculada(ws, 'B7', f'=IFERROR(MAX(J{d0}:J{d1}),"")', motor.FMT_PCT1, bold=True)

    # ---- tabla ---------------------------------------------------------
    for i, r in enumerate(range(d0, d1 + 1)):
        f = filas[i] if i < len(filas) else None
        entrada(ws, f'A{r}', f['ingrediente'] if f else None, align='left')
        entrada(ws, f'B{r}', f['categoria'] if f else None, align='left')
        entrada(ws, f'C{r}', None, align='left')
        entrada(ws, f'D{r}', f['formato'] if f else None, align='left')
        entrada(ws, f'E{r}', f['contenido'] if f else None, 'General')
        entrada(ws, f'F{r}', f['ub'] if f else None)
        entrada(ws, f'G{r}', f['precio_formato'] if f else None, motor.FMT_EUR)
        calculada(ws, f'H{r}',
                  f'=IF(OR(E{r}="",G{r}=""),"",IFERROR(ROUND(G{r}/E{r},4),""))',
                  FMT_EUR4, bold=True)
        entrada(ws, f'I{r}', f['anterior'] if f else None, FMT_EUR4)
        # ROUND a 6 decimales: una subida EXACTAMENTE igual al umbral (21/20-1)
        # sale 0,05000000000000004 en coma flotante y dispararía la alerta.
        calculada(ws, f'J{r}',
                  f'=IF(OR(H{r}="",I{r}=""),"",IFERROR(ROUND(H{r}/I{r}-1,6),""))',
                  motor.FMT_PCT1)
        calculada(ws, f'K{r}',
                  f'=IF(J{r}="","",IF(J{r}>$B$5,"ALERTA","OK"))', 'General', bold=True)
        entrada(ws, f'L{r}', None, FMT_FECHA)
        c = entrada(ws, f'M{r}', f['notas'] if f else None, align='left')
        c.font = Font(size=9, color=GRIS_NOTA)
        c.alignment = Alignment(horizontal='left', vertical='center', wrap_text=True)
        if f and f['notas']:
            # R1 CH-06: «Se usa en…» en todas las filas; sin altura, Excel deja la
            # fila a 15 pt y el texto ajustado se corta.
            lineas = -(-len(f['notas']) // CHARS_NOTA13)
            if lineas > 1:
                ws.row_dimensions[r].height = 12.0 * lineas + 3

    dv = motor._dv_lista('"' + ','.join(motor.UNIDADES) + '"',
                         'Elige una unidad de la lista.')
    ws.add_data_validation(dv)
    dv.add(f'F{d0}:F{d1}')
    ws.conditional_formatting.add(
        f'K{d0}:K{d1}', CellIsRule(operator='equal', formula=['"ALERTA"'],
                                   fill=_relleno(motor.ROJO_BG),
                                   font=Font(bold=True, color=motor.ROJO_FG)))
    ws.conditional_formatting.add(
        f'K{d0}:K{d1}', CellIsRule(operator='equal', formula=['"OK"'],
                                   fill=_relleno(VERDE_OK_BG),
                                   font=Font(bold=True, color=VERDE_OK_FG)))
    ws.auto_filter.ref = f'A{cab}:{ult}{d1}'
    return ws, d0, d1


def generar_13(destino, origen, datos, informe=None):
    informe = informe if informe is not None else []
    filas, conflictos, mapa, avisos = filas_lista(origen, datos)
    wb = openpyxl.Workbook()
    hoja_instrucciones(wb, 'Lista de Precios de Ingredientes',
                       instrucciones_13(len(filas), datos['lista_precios']['filas_vacias']))
    ws, d0, d1 = hoja_lista(wb, filas, datos)
    metadatos(wb, 'Lista de Precios de Ingredientes')
    _cerrar_comun(wb, F13, informe)
    # R1 CH-08: impresión sin Notas (M): del ~58 % al ~75 % de escala.
    ws.print_area = f'A1:{COL_IMPRESION13}{d1}'
    ws.print_title_rows = f'{LISTA["cab"]}:{LISTA["cab"]}'
    ws.freeze_panes = f'B{d0}'
    # R1 CH-02 (rebaja R2T-08): las columnas de fórmula H, J y K quedan
    # BLOQUEADAS (motor.proteger bloquea todo lo que no es verde). Con ellas
    # desbloqueadas, un Supr sobre los ejemplos —lo que pedían las Instrucciones—
    # borraba las fórmulas sin aviso. Excel no ordena un rango con celdas
    # bloqueadas aunque se dé el permiso, así que ORDENAR queda para la hoja
    # desprotegida (sin contraseña; lo dicen las Instrucciones) y el permiso de
    # ordenar se deja cerrado: así nadie ordena solo las columnas verdes y
    # separa el «Precio anterior» de su ingrediente. FILTRAR sí funciona
    # protegida (autofiltro permitido).
    ws.protection.sort = True
    ws.protection.autoFilter = False
    path = os.path.join(destino, F13)
    wb.save(path)
    return path, {'filas': len(filas), 'vacias': datos['lista_precios']['filas_vacias'],
                  'd0': d0, 'd1': d1, 'conflictos': conflictos, 'avisos': avisos,
                  'mapa': mapa, 'lista': filas}


# ==========================================================================
def merma_calculada(path12):
    """«MERMA % PARA TU ESCANDALLO» del 12, calculada por pycel (no a mano)."""
    import contextlib
    import logging
    logging.disable(logging.CRITICAL)
    from pycel import ExcelCompiler
    xl = ExcelCompiler(filename=path12)
    with open(os.devnull, 'w') as dn, contextlib.redirect_stderr(dn):
        v = xl.evaluate(f"'Test de despiece'!{DESP['merma']}")
        cp = xl.evaluate(f"'Test de despiece'!{DESP['coste_porcion']}")
    return float(v), float(cp)


if __name__ == '__main__':
    destino = sys.argv[1]
    origen = sys.argv[2] if len(sys.argv) > 2 else ORIGEN
    os.makedirs(destino, exist_ok=True)
    d = cargar_datos()
    print(generar_12(destino, d))
    p, info = generar_13(destino, origen, d)
    print(p, info['filas'], 'filas +', info['vacias'], 'vacías;',
          len(info['conflictos']), 'conflictos;', info['avisos'])
