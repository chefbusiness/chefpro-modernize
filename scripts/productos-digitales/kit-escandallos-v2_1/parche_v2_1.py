#!/usr/bin/env python3
"""
parche_v2_1.py — Parche v2.1 sobre los 12 libros publicados del Kit de Escandallos Pro.

Especificación: `recipe-costing-kit/SPEC.md` §2.0 («Arreglos del ES que la v2.1 incluye»),
R2T-15, R2T-16, R2-05 y R2-13. Trabaja SIEMPRE sobre la carpeta que le pasen (main.py le
pasa la copia del scratchpad salvo con `--real`). IDEMPOTENTE: todo son escrituras
absolutas o sustituciones que, aplicadas dos veces, dejan lo mismo.

  (a) 01-08: el rango `Conversiones!$A$5:$B$37` de las fórmulas del Factor pasa a
      `$A$5:$B${4+n+30}` (n = parejas de la tabla, hoy 33 → 67: tabla + 30 filas libres) y
      el área de impresión de «Conversiones» crece igual. La nota A2 de «Conversiones»
      dice hasta qué fila se pueden añadir parejas.
  (b) Los 12: línea «Versión 2.1 · <mes> 2026 · …» en Instrucciones y subject
      «Kit de Escandallos Pro · v2.1».
  (c) 01, fila 5 (solomillo): la merma H5 pasa a la «MERMA % PARA TU ESCANDALLO» que
      CALCULA el libro 12 (main.py la evalúa con pycel y la pasa aquí).
  (d) Texto, sin estructura nueva: una línea en las Instrucciones del 01 (junto a
      «Merma (%)») y en la hoja «Mermas» de 01-08 que remite al 12; una línea en las
      Instrucciones de 01-08 (junto a «Precio/Ud (€)») que remite al 13.
  Revisión R1 (INFORME §7):
  (c2) 08!Smash Burger!H5 (picada comprada lista) 15 % → 0 % [CH-03/T2].
  (c3) nombre y categoría de algunas filas de ejemplo [CH-10] (sólo texto).
  (c4) 03!Rotación Semanal: umbral del formato condicional en K3 [T6].

09, 10, 11 y BONUS se parchean a nivel de ZIP (core.xml + la celda inlineStr de la versión), no con
openpyxl: 09 y 11 llevan un gráfico y openpyxl los BORRA al cargar y guardar un libro.

NO inyecta caché: `inject_cache.py` va al final, en main.py.
"""
import copy
import datetime
import html
import os
import re
import sys
import zipfile

AQUI = os.path.dirname(os.path.abspath(__file__))
SCRIPTS = os.path.dirname(AQUI)
sys.path.insert(0, os.path.join(SCRIPTS, 'kit-escandallos-v2_0'))
sys.path.insert(0, AQUI)

import openpyxl                                                # noqa: E402
from openpyxl.styles import Alignment, Font, Protection        # noqa: E402

import motor                                                   # noqa: E402
import nuevas                                                  # noqa: E402

OTROS = ['09-control-mermas.xlsx', '10-calculadora-pvp.xlsx',
         '11-dashboard-food-cost-mensual.xlsx', 'BONUS-mermas-inventario.xlsx']
FILAS_LIBRES_CONV = 30
RX_CONV = re.compile(r"Conversiones!\$A\$5:\$B\$(\d+)")

CONV_A2_V20 = ('La usa la columna «Factor» de cada escandallo. Puedes añadir parejas '
               'al final: la clave se escribe «compra→uso».')


def conv_a2(fin):
    return ('La usa la columna «Factor» de cada escandallo. Puedes añadir parejas '
            f'al final, hasta la fila {fin}: la clave se escribe «compra→uso».')


MARCA_13 = '▸ ¿Llevas tus precios en la plantilla 13'
# R1 T3: un pegado normal NO da error: desplaza las referencias, la línea se queda
# en blanco sin avisar y el SUM del plato la ignora (coste más bajo).
LINEA_13 = ('▸ ¿Llevas tus precios en la plantilla 13 (Lista de Precios de Ingredientes)? '
            'Lleva su «Precio por unidad base» a «Precio/Ud (€)» con Pegado especial → '
            'Valores (en Excel, Ctrl+Alt+V y luego V) y pon esa misma unidad en «Ud. '
            'Compra». Un pegado normal arrastraría la fórmula: el precio y el coste de esa '
            'línea se quedarían en blanco sin avisar, y el coste del plato saldría más bajo.')
LINEA_13_04 = (' Los destilados, el vino y el espumoso no: su precio viene de la pestaña '
               '«Formatos de Compra».')
MARCA_12 = '▸ La merma del solomillo ('
# R1 T5: ≤ 110 caracteres, como A2 (114): la celda no ajusta texto y el área de
# impresión acaba en J, así que una línea más larga se cortaba en el papel.
MERMAS_A3 = ('Mide tu merma real con la plantilla 12 (Test de Rendimiento) y llévala a '
             '«Merma (%)» de la ficha.')
MERMAS_A3_PREVIAS = (
    'Mide tu propia merma con la plantilla 12 (Test de Rendimiento): pesa la '
    'pieza entera y lo que sale de ella, y lleva su «MERMA % PARA TU ESCANDALLO» '
    'a la columna «Merma (%)» de la ficha.',)

# R1 CH-03/T2: la picada de la smash burger se compra lista (merma de despiece 0 %)
# y la ficha escribe su peso CRUDO: el 15 % la encarecía un 18 % frente al test de
# cocción del 12, que usa esa misma ficha como ejemplo.
PICADA_08 = ('Carne picada (blend 80/20)', 'Carne picada (mezcla 80/20)')

# R1 T6: Google Sheets no admite referencias a otra hoja en el formato condicional.
ROT_CF_V20 = "AND(ISNUMBER(K5),K5>'Resumen Menú'!$C$13)"
ROT_CF_V21 = 'AND(ISNUMBER(K5),K5>$K$3)'
ROT_J3 = 'Food cost objetivo (del Resumen)'
ROT_K3 = "='Resumen Menú'!$C$13"


def linea_12(merma):
    generica = dict((m[0], m[1]) for m in motor.MERMAS)['Carne roja']
    return (f'▸ La merma del solomillo ({nuevas.pct(merma)}) no es la genérica de '
            f'«Carne roja» ({nuevas.pct(generica, 0)}): sale del test de despiece de la '
            'plantilla 12 (Test de Rendimiento), que ya descuenta lo que vale el cordón. '
            'Mide tu propia merma con la plantilla 12.')


# ==========================================================================
# Utilidades
# ==========================================================================
def insertar_tras(ws, prefijo_ancla, texto, marca, col=2):
    """Inserta `texto` en la fila siguiente a la que empieza por `prefijo_ancla`,
    bajando una fila lo que hay debajo (valores, estilos y alturas). Si ya existe una
    línea que empieza por `marca`, sólo reescribe su texto (idempotente).
    Las hojas de Instrucciones no tienen fórmulas, validaciones ni formatos
    condicionales, y su única combinación (B2) está por encima de cualquier ancla."""
    for r in range(1, ws.max_row + 1):
        v = ws.cell(row=r, column=col).value
        if isinstance(v, str) and v.startswith(marca):
            if v != texto:
                ws.cell(row=r, column=col).value = texto
                return 'reescrita'
            return 'ya estaba'
    ancla = None
    for r in range(1, ws.max_row + 1):
        v = ws.cell(row=r, column=col).value
        if isinstance(v, str) and v.startswith(prefijo_ancla):
            ancla = r
            break
    if ancla is None:
        raise SystemExit(f'{ws.title}: no encuentro la línea «{prefijo_ancla}…»')
    for m in ws.merged_cells.ranges:
        if m.max_row > ancla:
            raise SystemExit(f'{ws.title}: combinación {m} por debajo del ancla')
    ultima_col = max(ws.max_column, col)
    for r in range(ws.max_row, ancla, -1):
        for c in range(1, ultima_col + 1):
            src, dst = ws.cell(row=r, column=c), ws.cell(row=r + 1, column=c)
            dst.value = src.value
            dst._style = copy.copy(src._style)
        ws.row_dimensions[r + 1].height = ws.row_dimensions[r].height
    for c in range(1, ultima_col + 1):
        nueva = ws.cell(row=ancla + 1, column=c)
        nueva.value = None
        nueva._style = copy.copy(ws.cell(row=ancla, column=c)._style)
    ws.cell(row=ancla + 1, column=col).value = texto
    ws.row_dimensions[ancla + 1].height = ws.row_dimensions[ancla].height
    return f'insertada en fila {ancla + 1}'


def parejas_conversiones(ws):
    r = 5
    while ws.cell(row=r, column=1).value not in (None, ''):
        r += 1
    return r - 5


# ==========================================================================
# 01-08 (openpyxl: no llevan gráficos ni imágenes)
# ==========================================================================
def corregir_textos(wb, datos):
    """R1 CH-10: nombre y categoría de las filas de ejemplo de 01-08, con el mismo
    mapa que usa nuevas.leer_fichas() para el 13. Sólo texto: las mermas de las
    filas precargadas están escritas a mano, así que ningún coste cambia.
    Idempotente: una fila ya corregida no se toca."""
    hechos = []
    for ws, lay in motor.hojas_escandallo(wb):
        for r in range(lay.d0, lay.d1 + 1):
            a, b = ws.cell(row=r, column=1), ws.cell(row=r, column=2)
            if not (isinstance(a.value, str) and a.value.strip()):
                continue
            if isinstance(b.value, str) and b.value.startswith('='):
                continue
            nom, cat = nuevas.corregir_item(a.value.strip(), b.value, datos)
            if nom != a.value.strip():
                hechos.append(f'{ws.title}!A{r}: {a.value!r} → {nom!r}')
                a.value = nom
            if cat != b.value:
                hechos.append(f'{ws.title}!B{r}: {b.value!r} → {cat!r}')
                b.value = cat
    return hechos


def umbral_rotacion(ws):
    """R1 T6: el rojo de 03!'Rotación Semanal'!K5:K11 comparaba con otra hoja, algo
    que Google Sheets no admite en un formato condicional. El objetivo se trae a
    K3 (gris, bloqueada: viene del «Resumen Menú», como la columna «PVP menú») y la
    regla compara con $K$3. Idempotente."""
    if ws['K4'].value != 'FOOD COST (%)':
        raise SystemExit(f'03!Rotación Semanal!K4 = {ws["K4"].value!r}: no es el food cost')
    for coord, esperado in (('J3', ROT_J3), ('K3', ROT_K3)):
        if ws[coord].value not in (None, esperado):
            raise SystemExit(f'03!Rotación Semanal!{coord} ocupada: {ws[coord].value!r}')
    reglas = 0
    for rango in ws.conditional_formatting:
        for regla in rango.rules:
            if list(regla.formula or ()) in ([ROT_CF_V20], [ROT_CF_V21]):
                regla.formula = [ROT_CF_V21]
                reglas += 1
    if reglas != 1:
        raise SystemExit(f'03!Rotación Semanal: {reglas} reglas de food cost (se esperaba 1)')
    j3 = ws['J3']
    j3.value = ROT_J3
    j3.font = Font(bold=True, size=10)
    j3.alignment = Alignment(horizontal='right', vertical='center')
    j3.protection = Protection(locked=True)
    k3 = ws['K3']
    k3.value = ROT_K3
    k3._style = copy.copy(ws['J5']._style)           # gris «viene de otra hoja»
    k3.number_format = ws['K5'].number_format
    k3.protection = Protection(locked=True)
    return {'J3': ROT_J3, 'K3': ROT_K3, 'regla': ROT_CF_V21}


def parche_escandallo(path, fname, merma_solomillo, version, informe, datos=None):
    wb = openpyxl.load_workbook(path)
    rep = {'fichero': fname}

    # (a) rango de Conversiones ------------------------------------------
    conv = wb['Conversiones']
    n = parejas_conversiones(conv)
    fin = 4 + n + FILAS_LIBRES_CONV
    destino = f'Conversiones!$A$5:$B${fin}'
    total, cambiadas, raras = 0, 0, []
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                v = c.value
                if not (isinstance(v, str) and v.startswith('=') and 'Conversiones!' in v):
                    continue
                total += 1
                nuevo = RX_CONV.sub(destino, v)
                if RX_CONV.search(v) is None:
                    raras.append(f'{ws.title}!{c.coordinate}: {v[:80]}')
                if nuevo != v:
                    c.value = nuevo
                    cambiadas += 1
    conv.print_area = f'A1:J{fin}'
    a2 = conv['A2'].value
    if a2 in (CONV_A2_V20, conv_a2(fin)):
        conv['A2'].value = conv_a2(fin)
    else:
        informe.append(f'{fname}: Conversiones!A2 con texto inesperado, no se toca')
    rep.update({'parejas': n, 'rango': destino, 'formulas_conversiones': total,
                'formulas_reescritas': cambiadas, 'referencias_raras': raras,
                'area_conversiones': conv.print_area})

    # (b) versión y subject ----------------------------------------------
    motor.linea_instrucciones(wb['Instrucciones'], version, motor.RX_VERSION)
    wb.properties.subject = nuevas.SUBJECT

    # (c) 01: merma del solomillo desde el 12 -----------------------------
    if fname == motor.FICHEROS[0]:
        ws = wb['Escandallo']
        if ws['A5'].value != 'Solomillo de ternera':
            raise SystemExit(f'01!Escandallo!A5 = {ws["A5"].value!r}: no es el solomillo')
        antes = ws['H5'].value
        ws['H5'].value = round(merma_solomillo, 4)
        rep['H5'] = {'antes': antes, 'despues': ws['H5'].value}

    # (c2) 08: picada de la smash burger a 0 % (R1 CH-03/T2) ---------------
    if fname.startswith('08'):
        ws = wb['Smash Burger']
        if ws['A5'].value not in PICADA_08:
            raise SystemExit(f'08!Smash Burger!A5 = {ws["A5"].value!r}: no es la picada')
        antes = ws['H5'].value
        ws['H5'].value = 0
        rep['H5_08'] = {'antes': antes, 'despues': 0}

    # (c3) correcciones de texto de las filas de ejemplo (R1 CH-10) ---------
    rep['correcciones'] = corregir_textos(wb, datos or nuevas.cargar_datos())

    # (c4) 03: umbral del formato condicional en la misma hoja (R1 T6) -------
    if fname.startswith('03'):
        rep['rotacion'] = umbral_rotacion(wb['Rotación Semanal'])

    # (d) líneas de texto --------------------------------------------------
    ins = wb['Instrucciones']
    linea13 = LINEA_13 + (LINEA_13_04 if fname.startswith('04') else '')
    rep['linea_13'] = insertar_tras(ins, '▸ Precio/Ud (€):', linea13, MARCA_13)
    if fname == motor.FICHEROS[0]:
        rep['linea_12'] = insertar_tras(ins, '▸ Merma (%):', linea_12(merma_solomillo),
                                        MARCA_12)
    me = wb['Mermas']
    if me['A3'].value not in (None, MERMAS_A3) + MERMAS_A3_PREVIAS:
        raise SystemExit(f'{fname}: Mermas!A3 ocupada: {me["A3"].value!r}')
    me['A3'].value = MERMAS_A3
    me['A3']._style = copy.copy(me['A2']._style)
    rep['mermas_A3'] = 'ok'

    wb.save(path)
    informe.append(rep)
    return rep


# ==========================================================================
# 09, 10, 11, BONUS (ZIP: conserva los gráficos)
# ==========================================================================
RX_IS = re.compile(r'<is>(.*?)</is>', re.S)          # openpyxl 3.1: t="inlineStr"
RX_SI = re.compile(r'<si>(.*?)</si>', re.S)          # por si un libro trae sharedStrings
RX_T = re.compile(r'<t(\s[^>]*)?>(.*?)</t>', re.S)


def _texto(xml):
    return ''.join(html.unescape(m.group(2)) for m in RX_T.finditer(xml))


def parche_zip(path, fname, version, informe):
    with zipfile.ZipFile(path) as z:
        infos = z.infolist()
        partes = {i.filename: z.read(i.filename) for i in infos}
    rep = {'fichero': fname}

    core = partes['docProps/core.xml'].decode('utf-8')
    subj = html.escape(nuevas.SUBJECT, quote=False)
    core2, n_sub = re.subn(r'(<dc:subject[^>]*>)(.*?)(</dc:subject>)',
                           lambda m: m.group(1) + subj + m.group(3), core, flags=re.S)
    if n_sub != 1:
        raise SystemExit(f'{fname}: core.xml sin <dc:subject> único ({n_sub})')
    if core2 != core:
        ahora = datetime.datetime.utcnow().replace(microsecond=0).isoformat() + 'Z'
        core2 = re.sub(r'(<dcterms:modified[^>]*>)(.*?)(</dcterms:modified>)',
                       lambda m: m.group(1) + ahora + m.group(3), core2, flags=re.S)
    partes['docProps/core.xml'] = core2.encode('utf-8')

    hits = []
    nueva = html.escape(version, quote=False)

    def cambia_is(m):
        texto = _texto(m.group(1))
        if motor.RX_VERSION.match(texto):
            hits.append(texto)
            return '<is><t>' + nueva + '</t></is>'
        return m.group(0)

    def cambia_si(m):
        texto = _texto(m.group(1))
        if motor.RX_VERSION.match(texto):
            hits.append(texto)
            return '<si><t>' + nueva + '</t></si>'
        return m.group(0)

    for nombre in list(partes):
        if nombre.startswith('xl/worksheets/sheet') and nombre.endswith('.xml'):
            xml = partes[nombre].decode('utf-8')
            xml2 = RX_IS.sub(cambia_is, xml)
            if xml2 != xml:
                partes[nombre] = xml2.encode('utf-8')
                rep['hoja_version'] = nombre
        elif nombre == 'xl/sharedStrings.xml':
            xml = partes[nombre].decode('utf-8')
            partes[nombre] = RX_SI.sub(cambia_si, xml).encode('utf-8')
    if len(hits) != 1:
        raise SystemExit(f'{fname}: {len(hits)} líneas de versión (se esperaba 1)')
    rep['version_antes'] = hits[0]

    tmp = path + '.tmp'
    with zipfile.ZipFile(tmp, 'w', zipfile.ZIP_DEFLATED) as zout:
        for i in infos:
            zout.writestr(i, partes[i.filename])
    os.replace(tmp, path)
    rep['partes'] = len(infos)
    informe.append(rep)
    return rep


def aplicar(carpeta, merma_solomillo, informe=None, version=None, datos=None):
    informe = informe if informe is not None else []
    version = version or nuevas.version_line()
    datos = datos or nuevas.cargar_datos()
    for fname in motor.FICHEROS:
        parche_escandallo(os.path.join(carpeta, fname), fname, merma_solomillo,
                          version, informe, datos)
    for fname in OTROS:
        parche_zip(os.path.join(carpeta, fname), fname, version, informe)
    return informe


if __name__ == '__main__':
    import json
    print(json.dumps(aplicar(sys.argv[1], float(sys.argv[2])), ensure_ascii=False,
                     indent=1, default=str))
