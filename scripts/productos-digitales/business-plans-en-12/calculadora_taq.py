#!/usr/bin/env python3
"""
calculadora_taq.py — Taco Cost & Trompo Yield Calculator, D145 del Taco Shop Business Plan Kit (EN, $39). Es la edición EE. UU. de
`../plan-taqueria/calculadora_taq_es.py`: MISMAS hojas (`Trompo Yield` · `Tortillas` · `Taco Costing` · `Instructions`), filas,
fórmulas, helpers, `inject_cache` y autotest (pycel + caché = `calcular()`, 3 juegos). Cambios EN:

  (1) carne en libras ($/lb) y porciones en ONZAS, con la constante gris «16 oz per lb» (el gate prohíbe números tecleados en
      fórmulas salvo 0/1);
  (2) hoja `Tortillas`: tortilla de maíz de 6 in (Mission, 360/caja) frente a masa harina + máquina + 0.75 FTE de tortillera;
      el equilibrio en tortillas al día se calcula con la misma fórmula que el ES;
  (3) `Taco Costing`: street taco de DOBLE tortilla de 6 in a $4.00, sales tax en carta = 0 %, comisión de plataforma 25 %;
  (4) campo verde «Your county's al pastor cook temp (°F)» VACÍO: la temperatura la pone el condado (nunca una cifra única);
  (5) US Letter, formatos sin símbolo de moneda (los rótulos llevan «($)»), docProps en inglés, versión del libro EN.

Valores que otros scripts leen de la CACHÉ del xlsx (dict `EXPUESTOS`, función `valores_expuestos(ruta)`).
Corre en el VPS (openpyxl + pycel):

    /root/venv-guias/bin/python calculadora_taq.py --salida <dir>             # genera + inject_cache + autotest
    /root/venv-guias/bin/python calculadora_taq.py --salida <dir> --autotest

Sesión Claude Code, 29-oct-2026.
"""
import argparse
import math
import os
import re
import sys
from collections import OrderedDict

import openpyxl
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill, Protection
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

AQUI = os.path.dirname(os.path.abspath(__file__))
PD = os.path.dirname(AQUI)
os.environ.setdefault('BP_CONJUNTO', 'taq')
sys.path.insert(0, os.path.join(PD, 'planes-v2_0'))
sys.path.insert(0, os.path.join(PD, 'business-plans-en'))
sys.path.insert(0, PD)
sys.path.insert(0, AQUI)
import motor                                                                       # noqa: E402  (solo gate_formatos)
import mapas as M                                                                  # noqa: E402  (detectores y textos fijos)

_PL = M.PLANES['taq']
SLUG = _PL['slug']
PRODUCTO = _PL['producto']
URL = _PL['url']
FICHERO = 'taco-cost-trompo-yield-calculator.xlsx'
TITULO = 'Taco Cost & Trompo Yield Calculator (Al Pastor Cone, Tortillas and Taco Costing)'
KEYWORDS = _PL['keywords']
VERSION, MES = M.VERSION, M.MES_VERSION
VERSION_LINEA = 'Version %s · %s · %s · info@aichef.pro' % (VERSION, MES, URL)
FOOTER = M.FOOTER_EN
HT, HO, HE, HI = 'Trompo Yield', 'Tortillas', 'Taco Costing', 'Instructions'
VERDE = PatternFill('solid', fgColor='E8F5E9')
GRIS = PatternFill('solid', fgColor='F2F2F2')
ROJO = PatternFill('solid', bgColor='F8CBAD', fgColor='F8CBAD')
AMBAR = PatternFill('solid', bgColor='FFE699', fgColor='FFE699')
OK = PatternFill('solid', bgColor='C6E0B4', fgColor='C6E0B4')
FUENTE = 'Calibri'
USD, ENT, PCT, DEC, OZF, USD3 = '#,##0.00', '#,##0', '0.0%', '#,##0.0', '#,##0.00', '#,##0.000'
MODOS = ['Buy corn tortillas', 'Make them with masa harina']
FC_OK, FC_VIG, FC_ALTO = 'On target', 'Watch the cost', 'Above target'
SEM_COMPRA, SEM_PROPIA = 'Buy them', 'Make them'
EST = ' Kit estimate, editable.'
KIT_PRECIO = "Kit estimate (Oct 2026): replace with your distributor's price"

# ---------------------------------------------------------------------------------------------------------------------
# Datos (F1-research-us-delta §4-§5). Importes antes de sales tax. Todo «Kit estimate (Oct 2026)», editable.
# ---------------------------------------------------------------------------------------------------------------------
#: hoja Trompo Yield, bloque A: (clave, rótulo, formato, valor por defecto, constante gris, nota)
ENT_A = [
    ('K', 'Raw meat stacked on the cone (lb)', DEC, 30.0, False, 'Weigh the sliced meat before the marinade, ready to stack.' + EST),
    ('a', 'Marinade over the meat (%)', PCT, 0.08, False, 'Achiote, chiles, spices and pineapple in the marinade, over the weight of the meat.' + EST),
    ('Pc', 'Raw meat price ($/lb, before tax)', USD, 3.00, False, 'Pork shoulder. Type your supplier\'s price.' + EST),
    ('Pa', 'Marinade price ($/lb, before tax)', USD, 3.00, False, 'Type your supplier\'s price or the cost of your own recipe.' + EST),
    ('pk', 'Pineapple on top of the cone (lb)', DEC, 0.0, False, 'Default 0: the pineapple on each taco is counted separately on the "Taco Costing" sheet. If you type pounds here, they add to the cone cost.'),
    ('pp', 'Pineapple price ($/lb, before tax)', USD, 0.80, False, 'Only counts if you type pounds of pineapple above.' + EST),
    ('mc', 'Cooking and drip loss (%)', PCT, 0.35, False, 'What the cone loses as it roasts: fat and water. There is no reliable published figure: measure it with your first cone (see the "Instructions" sheet).' + EST),
    ('mr', 'Trim you cannot serve (%)', PCT, 0.05, False, 'Burnt tips, hard fat and gristle that never reach the plate.' + EST),
    ('S', 'Leftover at close (lb)', DEC, 0.0, False, 'Cooked meat left unserved at close. Weigh it; default 0.'),
    ('r', 'Share of the leftover you reuse (%)', PCT, 0.0, False, 'For example in a next-day stew. The part you throw away takes tacos off the total.'),
    ('gt', 'Cooked meat served per taco (oz)', OZF, 1.5, False, 'Ounces of cooked, sliced meat on each taco.' + EST),
    ('N', 'Number of tacos you want to prepare', ENT, 100, False, 'Whole number greater than 0. It tells you how much raw meat you need.'),
    ('OZ', 'Ounces per lb (constant: 16 oz per lb)', ENT, 16, True, 'Constant: 1 lb = 16 oz. Not editable. It converts the $/lb prices to ounces.'),
    ('ref', 'Benchmark: tacos per raw lb at a Mexico City stand', ENT, 11, True,
     'Fixed benchmark: a 110 to 220 lb trompo yields 1,250 to 2,500 tacos, about 11 per raw lb (Milenio). Not editable.'),
    ('ct', "Your county's al pastor cook temp (°F)", '0', None, False,
     'Leave it empty until you have it. Ask your county health department IN WRITING what cook and hold temperature it requires for '
     'al pastor off the trompo; counties differ, so the kit gives no single figure. Write it here and keep the answer on file.'),
]
SAL_A = [
    ('marin', 'Marinated meat (lb)', DEC, '{K}*(1+{a})', 'Raw meat plus marinade.'),
    ('servidos', 'Cooked meat served from the cone (lb)', '#,##0.00', '{marin}*(1-{mc}-{mr})-{S}*(1-{r})',
     'Marinated meat minus the loss and the trim, minus the leftover you throw away.'),
    ('rend', 'TROMPO YIELD (%)', PCT, '{servidos}/{K}', 'Served meat ÷ raw meat stacked.'),
    ('tacos', 'TACOS FROM THE CONE', ENT, '{servidos}*{OZ}/{gt}', 'Served lb × 16 ÷ ounces per taco.'),
    ('tacos_lb', 'Tacos per raw lb', DEC, '{tacos}/{K}', 'Compare it with the Mexico City benchmark above.'),
    ('vs_ref', 'Your tacos per lb against the Mexico City benchmark (%)', '0%', '{tacos_lb}/{ref}', 'Below 100 %: you serve more ounces per taco or lose more.'),
    ('coste_cono', 'Cone cost ($)', USD, '{K}*{Pc}+{K}*{a}*{Pa}+{pk}*{pp}', 'Meat + marinade + pineapple on top.'),
    ('coste_kg', 'COST PER SERVED LB ($)', USD, '{coste_cono}/{servidos}', 'Cone cost ÷ served lb. It feeds the taco costing.'),
    ('carne_taco', 'MEAT COST PER TACO ($)', USD3, '{coste_kg}*{gt}/{OZ}', 'Cost per served lb × ounces per taco ÷ 16.'),
    ('kgN', 'Raw meat to prepare the N tacos (lb)', '#,##0.00', '{N}*{gt}/({OZ}*{rend})', 'N × ounces per taco ÷ (16 × yield).'),
]
#: bloque B: (guiso, lb crudas, $/lb crudo, rendimiento tras la cocción, oz por taco)
GUISOS = [
    ('Carnitas', 20.0, 2.80, 0.60, 1.5),
    ('Suadero (beef navel)', 20.0, 4.20, 0.55, 1.5),
    ('Cochinita pibil', 20.0, 3.00, 0.60, 1.75),
    ('Beef barbacoa', 20.0, 4.00, 0.55, 1.75),
    ('Chicken tinga', 20.0, 2.20, 0.65, 1.5),
    ('Campechano (steak and chorizo)', 20.0, 4.50, 0.60, 1.5),
]
NB = len(GUISOS)

#: hoja Tortillas: (clave, rótulo, formato, valor, nota)
ENT_T = [
    ('tp', 'Price of a case of corn tortillas ($, before tax)', USD, 58.99, '6 in white corn tortillas, 360 per case (Mission, case price read in October 2026).' + EST),
    ('tn', 'Tortillas per case', ENT, 360, 'Whole number greater than 0.' + EST),
    ('tg', 'Weight of each tortilla (oz)', OZF, 0.9, 'A 6 in corn tortilla weighs about 0.9 oz.' + EST),
    ('mz', 'Masa harina price ($/lb, before tax)', USD, 0.45, 'Instant corn masa flour in a 50 lb bag. Type your supplier\'s price.' + EST),
    ('mm', 'Masa you get from 1 lb of masa harina (lb)', '#,##0.00', 1.8, 'Masa harina absorbs water: 1 lb of flour makes about 1.8 lb of masa.' + EST),
    ('cg', 'Water, power and gas for the press ($ per lb of masa harina)', USD, 0.20, 'Everything you spend to mix, press and cook, per pound of masa harina.' + EST),
    ('gn', 'Raw masa per homemade tortilla (oz)', OZF, 1.0, 'Raw masa in each tortilla before pressing.' + EST),
    ('tw', 'Yearly cost of the tortilla maker ($, with payroll costs)', USD, 27450.0, 'About 0.75 FTE at $16 an hour with payroll costs.' + EST),
    ('ti', 'Tortilla machine investment ($, before tax)', USD, 7808.0, 'A commercial countertop tortilla machine with a 6 in cutter.' + EST),
    ('ty', 'Years of life of the machine', ENT, 7, 'The investment is spread over these years.' + EST),
    ('td', 'Days open per year', ENT, 308, 'The days of the kit\'s business plan case.' + EST),
    ('tpd', 'Tortillas you use per day', ENT, 650, 'The kit\'s case uses about 650 a day (tacos with two tortillas, plus quesadillas). Replace it with yours.' + EST),
]
SAL_T = [
    ('c_comp', 'Cost per bought tortilla ($)', USD3, '{tp}/{tn}', 'Case price ÷ tortillas per case.'),
    ('tort_kg', 'Tortillas in 1 lb bought', DEC, '{OZ}/{tg}', 'To order from your supplier: 16 ÷ ounces per tortilla.'),
    ('maiz_t', 'Masa harina per homemade tortilla (oz)', OZF, '{gn}/{mm}', 'Raw masa per tortilla ÷ masa from 1 lb of flour.'),
    ('tort_maiz', 'Tortillas from 1 lb of masa harina', DEC, '{mm}*{OZ}/{gn}', ''),
    ('c_prop', 'Cost per homemade tortilla ($, without the person or the machine)', USD3, '{maiz_t}*({mz}+{cg})/{OZ}', 'Masa harina + water, power and gas.'),
    ('ahorro', 'Saving per tortilla if you make it ($)', USD3, '{c_comp}-{c_prop}', 'If it is zero or negative, buying always wins.'),
    ('fijo_dia', 'Fixed daily cost of making them ($)', USD, '({tw}+{ti}/{ty})/{td}', 'Person + machine spread over the days you are open.'),
    ('eq', 'BREAK-EVEN (tortillas per day)', ENT, 'IF({ahorro}>0,({tw}+{ti}/{ty})/({ahorro}*{td}),0)',
     'Above this many tortillas a day, making them pays. Below it, buy.'),
    ('dia_comp', 'Daily cost buying them ($)', USD, '{tpd}*{c_comp}', 'For the tortillas per day you typed.'),
    ('dia_prop', 'Daily cost making them ($)', USD, '{tpd}*{c_prop}+{fijo_dia}', 'Includes the person and the machine.'),
    ('sem', 'What suits you', 'General', 'IF(AND({ahorro}>0,{tpd}>={eq}),"%s","%s")' % (SEM_PROPIA, SEM_COMPRA), 'Your tortillas per day against the break-even.'),
]

#: hoja Taco Costing
ESC_IN = [
    ('tax', 'Sales tax included in your menu prices (%)', PCT, 0.0, 'Typically 0 %: a US menu is quoted before tax and the tax is added at the register. Rates vary by state and city; ask your accountant.'),
    ('obj', 'Target food cost (% of the price before tax)', PCT, 0.25, 'Taco shop reference: 22 to 30 % on the al pastor taco.' + EST),
    ('tol', 'Watch margin over the target (points)', PCT, 0.05, 'Up to this margin above the target, the traffic light says "Watch the cost".'),
    ('merm', 'Production waste over the cost (%)', PCT, 0.05, 'Raw salsas you discard, torn tortillas, onion and cilantro you cannot use.' + EST),
    ('modo', 'Corn tortilla the costing uses', 'General', MODOS[0], 'Picks the cost per tortilla from the "Tortillas" sheet: bought or homemade with masa harina.'),
    ('env', 'Packaging per takeout order ($, before tax)', USD, 0.60, 'Bag, container and salsa cups of one order.' + EST),
    ('com', 'Delivery app commission (%)', PCT, 0.25, 'Over the price before tax. Check your contract with each app.' + EST),
]
#: onzas: (ingrediente, $/lb o enlace 'cono'/'bN', origen del precio)
ING_G = [
    ('Al pastor meat (from the cone)', 'cono', 'linked from the Trompo Yield sheet ($ per served lb)'),
    ('Carnitas', 'b0', 'linked from the Trompo Yield sheet ($ per cooked lb)'),
    ('Suadero (beef navel)', 'b1', 'linked from the Trompo Yield sheet ($ per cooked lb)'),
    ('Cochinita pibil', 'b2', 'linked from the Trompo Yield sheet ($ per cooked lb)'),
    ('Beef barbacoa', 'b3', 'linked from the Trompo Yield sheet ($ per cooked lb)'),
    ('Chicken tinga', 'b4', 'linked from the Trompo Yield sheet ($ per cooked lb)'),
    ('Campechano (steak and chorizo)', 'b5', 'linked from the Trompo Yield sheet ($ per cooked lb)'),
    ('Diced white onion', 0.80, KIT_PRECIO),
    ('Cilantro', 6.00, KIT_PRECIO),
    ('Pineapple', 0.80, KIT_PRECIO),
    ('Lime', 1.80, KIT_PRECIO),
    ('Red salsa', 2.00, KIT_PRECIO),
    ('Green salsa', 2.00, KIT_PRECIO),
    ('Salsa macha (chile oil salsa)', 8.00, KIT_PRECIO),
    ('Oaxaca or Monterey Jack cheese', 6.00, KIT_PRECIO),
    ('Guacamole', 6.00, KIT_PRECIO + ' (depends on the avocado)'),
    ('Pinto beans', 1.50, KIT_PRECIO),
]
#: unidades: (ingrediente, $/ud o None si es el de la hoja Tortillas, origen)
ING_U = [
    ('Corn tortilla, 6 in ($/ea)', None, 'linked from the Tortillas sheet, according to the option you pick'),
    ('Flour tortilla, 10 in ($/ea)', 0.25, KIT_PRECIO),
]
TM, TH = 'Corn tortilla, 6 in ($/ea)', 'Flour tortilla, 10 in ($/ea)'
PAST, CAR, SUA, COC, BAR, TIN, CAM = [g[0] for g in ING_G[:7]]
CEB, CIL, PIN, LIM, SRO, SVE, SMA, QUE, GUA, FRI = [g[0] for g in ING_G[7:]]
#: columna → (precio de carta, {ingrediente en oz: oz}, {tortillas: uds})
RECETAS = OrderedDict([
    ('Al pastor taco', (4.00, {PAST: 1.5, CEB: 0.35, CIL: 0.1, PIN: 0.3, LIM: 0.2, SRO: 0.3, SVE: 0.2}, {TM: 2})),
    ('Carnitas taco', (4.00, {CAR: 1.5, CEB: 0.35, CIL: 0.1, LIM: 0.2, SRO: 0.25, SVE: 0.25}, {TM: 2})),
    ('Suadero taco', (4.50, {SUA: 1.5, CEB: 0.35, CIL: 0.1, LIM: 0.2, SRO: 0.25, SMA: 0.15, SVE: 0.1}, {TM: 2})),
    ('Cochinita pibil taco', (4.50, {COC: 1.75, CEB: 0.35, LIM: 0.2, SVE: 0.35}, {TM: 2})),
    ('Beef barbacoa taco', (4.75, {BAR: 1.75, CEB: 0.35, CIL: 0.1, LIM: 0.2, SRO: 0.4}, {TM: 2})),
    ('Chicken tinga taco', (4.00, {TIN: 1.5, CEB: 0.2, CIL: 0.1, QUE: 0.35, SRO: 0.3, SVE: 0.2}, {TM: 2})),
    ('Campechano taco', (4.75, {CAM: 1.5, CEB: 0.35, CIL: 0.1, LIM: 0.2, SRO: 0.3, SVE: 0.2}, {TM: 2})),
    ('Veggie taco', (4.50, {FRI: 1.4, QUE: 0.7, GUA: 0.9, CEB: 0.2, CIL: 0.1, LIM: 0.2, SVE: 0.35}, {TM: 2})),
    ('Al pastor order (3 tacos)', (11.00, {PAST: 4.5, CEB: 1.05, CIL: 0.3, PIN: 0.9, LIM: 0.6, SRO: 0.9, SVE: 0.6}, {TM: 6})),
    ('Al pastor gringa', (11.50, {PAST: 2.5, QUE: 2.0, PIN: 0.5, CEB: 0.35, CIL: 0.1, SRO: 0.4, SVE: 0.3}, {TH: 1})),
])
NP, NG, NU = len(RECETAS), len(ING_G), len(ING_U)
PCOL = [get_column_letter(4 + j) for j in range(NP)]                              # D..M
J_PASTOR, J_ORDEN, J_GRINGA = 0, NP - 2, NP - 1

# ---------------------------------------------------------------------------------------------------------------------
# Disposición de las hojas (filas)
# ---------------------------------------------------------------------------------------------------------------------
M_ = OrderedDict()                                                                 # hoja Trompo, entradas A → fila
M_HA = 5
for _i, (_k, *_r) in enumerate(ENT_A):
    M_[_k] = M_HA + 1 + _i
M_RA = M_HA + 1 + len(ENT_A) + 1
S_A = OrderedDict((k, M_RA + 1 + i) for i, (k, *_r) in enumerate(SAL_A))
M_HB = M_RA + 1 + len(SAL_A) + 1
B0 = M_HB + 1                                                                      # primera fila de guisos
B_NOTA = B0 + NB

T = OrderedDict()                                                                  # hoja Tortillas
_f_ = 5
for _k, *_r in ENT_T:
    _f_ += 1
    if _k in ('mz', 'tpd'):
        _f_ += 1                                                                   # fila de cabecera de sección
    T[_k] = _f_
T_RES = _f_ + 2
S_T = OrderedDict((k, T_RES + 1 + i) for i, (k, *_r) in enumerate(SAL_T))

E = OrderedDict((k, 5 + i) for i, (k, *_r) in enumerate(ESC_IN))
E_CAB = 5 + len(ESC_IN) + 1
E_PVP = E_CAB + 1
E_G0 = E_PVP + 1
E_G1 = E_G0 + NG - 1
E_UH = E_G1 + 1
E_U0 = E_UH + 1
E_U1 = E_U0 + NU - 1
E_R = E_U1 + 2
RES = [
    ('peso', 'Weight served on the plate, without tortilla (oz)', OZF),
    ('c_mat', 'Raw material cost ($)', USD3),
    ('c_tot', 'COST WITH WASTE ($)', USD3),
    ('pvp_sin', 'Menu price before tax ($)', USD),
    ('fc', 'FOOD COST (%)', PCT),
    ('sem', 'Traffic light against the target', 'General'),
    ('margen', 'Gross margin ($, price before tax − cost)', USD),
    ('pvp_sug', 'Menu price with tax needed to hit the target ($)', USD),
]
R = OrderedDict((k, E_R + 1 + i) for i, (k, *_r) in enumerate(RES))
E_CH = E_R + len(RES) + 2
CANAL = [
    ('sala_e', 'Dine-in: margin ($)', USD), ('sala_p', 'Dine-in: margin (%)', PCT),
    ('llev_e', 'Takeout: margin ($, after packaging)', USD), ('llev_p', 'Takeout: margin (%)', PCT),
    ('plat_e', 'Delivery app: margin ($, after packaging and commission)', USD), ('plat_p', 'Delivery app: margin (%)', PCT),
]
CH = OrderedDict((k, E_CH + 1 + i) for i, (k, *_r) in enumerate(CANAL))
CH_COL = {J_ORDEN: 'D', J_GRINGA: 'E'}                                             # columna del bloque de canal → receta
E_FIN = E_CH + len(CANAL) + 2

EXPUESTOS = OrderedDict([
    ('TACOS_POR_CONO', (HT, 'B%d' % S_A['tacos'])),
    ('RENDIMIENTO_TROMPO', (HT, 'B%d' % S_A['rend'])),
    ('COSTE_TACO_PASTOR', (HE, '%s%d' % (PCOL[J_PASTOR], R['c_tot']))),
    ('FC_TACO_PASTOR', (HE, '%s%d' % (PCOL[J_PASTOR], R['fc']))),
    ('EQUILIBRIO_TORTILLAS_DIA', (HO, 'B%d' % S_T['eq'])),
])


def valores_defecto():
    """Todas las entradas por defecto en un solo dict plano (claves = las de las celdas verdes)."""
    v = OrderedDict()
    for k, _r, _f2, val, _c, _n in ENT_A:
        v[k] = val
    for i, (_n, kg, p, rd, g) in enumerate(GUISOS):
        v['bk%d' % i], v['bp%d' % i], v['br%d' % i], v['bg%d' % i] = kg, p, rd, g
    for k, _r, _f2, val, _n in ENT_T:
        v[k] = val
    for k, _r, _f2, val, _n in ESC_IN:
        v[k] = val
    for i, (_n, p, _o) in enumerate(ING_G):
        if not isinstance(p, str):
            v['p%d' % i] = p
    v['th'] = ING_U[1][1]
    for j, (_nom, (pvp, gr, ud)) in enumerate(RECETAS.items()):
        v['pvp%d' % j] = pvp
        for i, (ni, *_r) in enumerate(ING_G):
            v['g%d_%d' % (i, j)] = float(gr.get(ni, 0))
        for i, (ni, *_r) in enumerate(ING_U):
            v['u%d_%d' % (i, j)] = float(ud.get(ni, 0))
    return v


# ---------------------------------------------------------------------------------------------------------------------
# Cálculo en Python (espejo de las fórmulas) — fuente del autotest
# ---------------------------------------------------------------------------------------------------------------------
def calcular(v):
    r = OrderedDict()
    r['marin'] = v['K'] * (1 + v['a'])
    r['servidos'] = r['marin'] * (1 - v['mc'] - v['mr']) - v['S'] * (1 - v['r'])
    r['rend'] = r['servidos'] / v['K']
    r['tacos'] = r['servidos'] * v['OZ'] / v['gt']
    r['tacos_lb'] = r['tacos'] / v['K']
    r['vs_ref'] = r['tacos_lb'] / v['ref']
    r['coste_cono'] = v['K'] * v['Pc'] + v['K'] * v['a'] * v['Pa'] + v['pk'] * v['pp']
    r['coste_kg'] = r['coste_cono'] / r['servidos']
    r['carne_taco'] = r['coste_kg'] * v['gt'] / v['OZ']
    r['kgN'] = v['N'] * v['gt'] / (v['OZ'] * r['rend'])
    r['B'] = []
    for i in range(NB):
        b = OrderedDict()
        b['coc'] = v['bk%d' % i] * v['br%d' % i]
        b['ckg'] = v['bp%d' % i] / v['br%d' % i]
        b['tacos'] = b['coc'] * v['OZ'] / v['bg%d' % i]
        b['ctaco'] = b['ckg'] * v['bg%d' % i] / v['OZ']
        r['B'].append(b)
    t = OrderedDict()
    t['c_comp'] = v['tp'] / v['tn']
    t['tort_kg'] = v['OZ'] / v['tg']
    t['maiz_t'] = v['gn'] / v['mm']
    t['tort_maiz'] = v['mm'] * v['OZ'] / v['gn']
    t['c_prop'] = t['maiz_t'] * (v['mz'] + v['cg']) / v['OZ']
    t['ahorro'] = t['c_comp'] - t['c_prop']
    t['fijo_dia'] = (v['tw'] + v['ti'] / v['ty']) / v['td']
    t['eq'] = (v['tw'] + v['ti'] / v['ty']) / (t['ahorro'] * v['td']) if t['ahorro'] > 0 else 0
    t['dia_comp'] = v['tpd'] * t['c_comp']
    t['dia_prop'] = v['tpd'] * t['c_prop'] + t['fijo_dia']
    t['sem'] = SEM_PROPIA if (t['ahorro'] > 0 and v['tpd'] >= t['eq']) else SEM_COMPRA
    r['T'] = t
    pg = []
    for i, (_n, p, _o) in enumerate(ING_G):
        pg.append(r['coste_kg'] if p == 'cono' else r['B'][int(p[1:])]['ckg'] if isinstance(p, str) else v['p%d' % i])
    r['pg'] = pg
    r['pu'] = [t['c_comp'] if v['modo'] == MODOS[0] else t['c_prop'], v['th']]
    for j in range(NP):
        gs = [v['g%d_%d' % (i, j)] for i in range(NG)]
        us = [v['u%d_%d' % (i, j)] for i in range(NU)]
        p = OrderedDict()
        p['peso'] = sum(gs)
        p['c_mat'] = sum(g_ * q for g_, q in zip(gs, pg)) / v['OZ'] + sum(u_ * q for u_, q in zip(us, r['pu']))
        p['c_tot'] = p['c_mat'] * (1 + v['merm'])
        p['pvp_sin'] = v['pvp%d' % j] / (1 + v['tax'])
        p['fc'] = p['c_tot'] / p['pvp_sin']
        p['sem'] = FC_OK if p['fc'] <= v['obj'] else (FC_VIG if p['fc'] <= v['obj'] + v['tol'] else FC_ALTO)
        p['margen'] = p['pvp_sin'] - p['c_tot']
        p['pvp_sug'] = p['c_tot'] / v['obj'] * (1 + v['tax'])
        p['sala_e'] = p['margen']
        p['sala_p'] = p['sala_e'] / p['pvp_sin']
        p['llev_e'] = p['sala_e'] - v['env']
        p['llev_p'] = p['llev_e'] / p['pvp_sin']
        p['plat_e'] = p['llev_e'] - v['com'] * p['pvp_sin']
        p['plat_p'] = p['plat_e'] / p['pvp_sin']
        r['p%d' % j] = p
    return r


# ---------------------------------------------------------------------------------------------------------------------
# Helpers (de la Pizza EN)
# ---------------------------------------------------------------------------------------------------------------------
def _f(ws, coord, v, b=False, sz=11, i=False, fmt=None, wrap=False, fill=None, locked=True):
    c = ws[coord]
    c.value = v
    c.font = Font(name=FUENTE, size=sz, bold=b, italic=i)
    if fmt:
        c.number_format = fmt
    c.alignment = Alignment(wrap_text=wrap, vertical='top')
    if fill:
        c.fill = fill
    c.protection = Protection(locked=locked)
    return c


def _nota(ws, fila, texto, desde='E', hasta='F', chars=80.0):
    if not texto:
        return
    _f(ws, '%s%d' % (desde, fila), texto, wrap=True)
    if desde != hasta:
        ws.merge_cells('%s%d:%s%d' % (desde, fila, hasta, fila))
    ws.row_dimensions[fila].height = max(16, 15 * math.ceil(len(texto) / chars))


def _imprimir(ws, horizontal=False):
    ws.page_setup.paperSize = M.PAPEL_EN                                          # 1 = US Letter
    ws.page_setup.orientation = 'landscape' if horizontal else 'portrait'
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.oddFooter.center.text = FOOTER
    ws.print_options.horizontalCentered = True
    ws.protection.sheet = True                                                    # sin contraseña
    ws.sheet_view.showGridLines = False


def _dv(ws, tipo, celdas, f1=None, f2=None, op=None, titulo='Invalid value', msg='', blank=False):
    dv = DataValidation(type=tipo, formula1=f1, formula2=f2, operator=op, allow_blank=blank, showErrorMessage=True,
                        errorTitle=titulo, error=msg)
    ws.add_data_validation(dv)
    for c in celdas:
        dv.add(c)


def a_excel(expr, ref):
    return re.sub(r'\{(\w+)\}', lambda m: ref[m.group(1)], expr)


def _oz_ext():
    return "'%s'!$B$%d" % (HT, M_['OZ'])


# ---------------------------------------------------------------------------------------------------------------------
# Hojas
# ---------------------------------------------------------------------------------------------------------------------
def hoja_trompo(wb):
    ws = wb.active
    ws.title = HT
    for col, w in (('A', 54), ('B', 15), ('C', 15), ('D', 16), ('E', 14), ('F', 14), ('G', 14), ('H', 12), ('I', 15)):
        ws.column_dimensions[col].width = w
    _f(ws, 'A1', 'TROMPO YIELD CALCULATOR: HOW MANY TACOS COME OUT OF THE CONE AND WHAT THE MEAT COSTS', b=True, sz=14)
    _f(ws, 'A2', 'Type only in the GREEN cells: everything else is calculated. Amounts before sales tax. Every default value is a kit '
                 'estimate, editable: replace it with your supplier\'s numbers and, above all, with what you measure on your own cone.', wrap=True)
    ws.merge_cells('A2:I2')
    ws.row_dimensions[2].height = 32
    _f(ws, 'A3', 'FIXED NOTICE. The cooking and drip loss (35 %) has no reliable published source: it is a kit estimate. Weigh your first '
                 'raw cone, what you serve and the leftover, and adjust the loss until the yield matches what you measure.',
       wrap=True, b=True, fill=GRIS)
    ws.merge_cells('A3:I3')
    ws.row_dimensions[3].height = 36
    _f(ws, 'A%d' % M_HA, 'A. THE AL PASTOR CONE', b=True, fill=GRIS)
    _f(ws, 'B%d' % M_HA, 'Value', b=True, fill=GRIS)
    _f(ws, 'C%d' % M_HA, 'Notes', b=True, fill=GRIS)
    for k, rot, fmt, val, const, nota in ENT_A:
        f = M_[k]
        _f(ws, 'A%d' % f, rot)
        if const:
            _f(ws, 'B%d' % f, val, fmt=fmt, fill=GRIS)
        else:
            _f(ws, 'B%d' % f, val, fmt=fmt, fill=VERDE, locked=False)
        _nota(ws, f, nota, desde='C', hasta='I', chars=105.0)
    ref = dict((k, '$B$%d' % f) for k, f in M_.items())
    ref.update((k, '$B$%d' % f) for k, f in S_A.items())
    for col, t in zip('ABC', ('CONE RESULT', 'Value', 'Notes')):
        _f(ws, '%s%d' % (col, M_RA), t, b=True, fill=GRIS)
    for k, rot, fmt, expr, nota in SAL_A:
        f = S_A[k]
        neg = k in ('rend', 'tacos', 'coste_kg', 'carne_taco')
        _f(ws, 'A%d' % f, rot, b=neg)
        _f(ws, 'B%d' % f, '=' + a_excel(expr, ref), fmt=fmt, b=neg)
        _nota(ws, f, nota, desde='C', hasta='I', chars=105.0)
    cab = ('B. OTHER PROTEINS (stews and roasts)', 'Raw lb', '$/raw lb (before tax)', 'Yield after cooking (%)', 'Oz per taco',
           'Cooked lb', '$/cooked lb', 'Tacos', '$ of meat per taco')
    for i, t in enumerate(cab):
        _f(ws, '%s%d' % (get_column_letter(1 + i), M_HB), t, b=True, fill=GRIS, wrap=True)
    ws.row_dimensions[M_HB].height = 46
    for i, (nom, kg, p, rd, g) in enumerate(GUISOS):
        f = B0 + i
        _f(ws, 'A%d' % f, nom, fill=VERDE, locked=False)
        _f(ws, 'B%d' % f, kg, fmt=DEC, fill=VERDE, locked=False)
        _f(ws, 'C%d' % f, p, fmt=USD, fill=VERDE, locked=False)
        _f(ws, 'D%d' % f, rd, fmt='0%', fill=VERDE, locked=False)
        _f(ws, 'E%d' % f, g, fmt=OZF, fill=VERDE, locked=False)
        _f(ws, 'F%d' % f, '=B%d*D%d' % (f, f), fmt='#,##0.00')
        _f(ws, 'G%d' % f, '=C%d/D%d' % (f, f), fmt=USD)
        _f(ws, 'H%d' % f, "=F%d*%s/E%d" % (f, '$B$%d' % M_['OZ'], f), fmt=ENT)
        _f(ws, 'I%d' % f, "=G%d*E%d/%s" % (f, f, '$B$%d' % M_['OZ']), fmt=USD3, b=True)
    _f(ws, 'A%d' % B_NOTA, 'Slow-cooked stews (carnitas, barbacoa, cochinita) yield less than the cone: the yield and the price on each row are '
                           'a kit estimate, editable. The $/cooked lb of each row feeds the "%s" sheet.' % HE,
       i=True, wrap=True)
    ws.merge_cells('A%d:I%d' % (B_NOTA, B_NOTA))
    ws.row_dimensions[B_NOTA].height = 34
    cel = lambda k: 'B%d' % M_[k]
    _dv(ws, 'decimal', [cel(k) for k in ('K', 'Pc', 'Pa', 'gt')], f1='0', op='greaterThan', msg='Type a number greater than 0.')
    _dv(ws, 'decimal', [cel(k) for k in ('pk', 'pp', 'S')], f1='0', op='greaterThanOrEqual', msg='Type a number greater than or equal to 0.')
    _dv(ws, 'decimal', [cel(k) for k in ('a', 'r')], f1='0', f2='1', op='between', titulo='Invalid percentage',
        msg='Type a percentage between 0 and 1 (0.08 = 8 %).')
    _dv(ws, 'decimal', [cel('mc')], f1='0', f2='0.6', op='between', titulo='Invalid loss', msg='Type between 0 and 0.6 (0.35 = 35 %).')
    _dv(ws, 'decimal', [cel('mr')], f1='0', f2='0.3', op='between', titulo='Invalid trim', msg='Type between 0 and 0.3 (0.05 = 5 %).')
    _dv(ws, 'whole', [cel('N')], f1='1', f2='1000000', op='between', titulo='Invalid number', msg='Type a whole number of 1 or more.')
    _dv(ws, 'decimal', [cel('ct')], f1='100', f2='212', op='between', titulo='Invalid temperature',
        msg='Type the temperature in degrees Fahrenheit (100 to 212), or leave it empty.', blank=True)
    rows = range(B0, B0 + NB)
    _dv(ws, 'decimal', ['%s%d' % (c, f) for c in 'BCE' for f in rows], f1='0', op='greaterThan', msg='Type a number greater than 0.')
    _dv(ws, 'decimal', ['D%d' % f for f in rows], f1='0.05', f2='1', op='between', titulo='Invalid yield',
        msg='Type between 0.05 and 1 (0.6 = 60 %).')
    _imprimir(ws, horizontal=True)


def hoja_tortilla(wb):
    ws = wb.create_sheet(HO)
    for col, w in (('A', 56), ('B', 17), ('C', 70)):
        ws.column_dimensions[col].width = w
    _f(ws, 'A1', 'TORTILLAS: BUY THEM OR MAKE THEM WITH MASA HARINA, AND THE BREAK-EVEN IN TORTILLAS PER DAY', b=True, sz=14)
    _f(ws, 'A2', 'Type only in the GREEN cells. Amounts before sales tax. Every default value is a kit estimate, editable: replace it with '
                 'your supplier\'s numbers. At the volume of a neighborhood taco shop, buying usually wins; this sheet shows you where the math flips.',
       wrap=True)
    ws.merge_cells('A2:C2')
    ws.row_dimensions[2].height = 34
    ref = dict((k, '$B$%d' % f) for k, f in T.items())
    ref.update((k, '$B$%d' % f) for k, f in S_T.items())
    ref['OZ'] = _oz_ext()
    sec = {'tp': 'A. BOUGHT TORTILLAS', 'mz': 'B. HOMEMADE (masa harina, tortilla machine and a tortilla maker)', 'tpd': 'C. YOUR CASE'}
    for k, rot, fmt, val, nota in ENT_T:
        f = T[k]
        if k in sec:
            _f(ws, 'A%d' % (f - 1), sec[k], b=True, fill=GRIS)
            _f(ws, 'B%d' % (f - 1), 'Value', b=True, fill=GRIS)
            _f(ws, 'C%d' % (f - 1), 'Notes', b=True, fill=GRIS)
        _f(ws, 'A%d' % f, rot)
        _f(ws, 'B%d' % f, val, fmt=fmt, fill=VERDE, locked=False)
        _nota(ws, f, nota, desde='C', hasta='C', chars=78.0)
    for col, t in zip('ABC', ('RESULT', 'Value', 'Notes')):
        _f(ws, '%s%d' % (col, T_RES), t, b=True, fill=GRIS)
    for k, rot, fmt, expr, nota in SAL_T:
        f = S_T[k]
        neg = k in ('c_comp', 'c_prop', 'eq', 'sem')
        _f(ws, 'A%d' % f, rot, b=neg)
        _f(ws, 'B%d' % f, '=' + a_excel(expr, ref), fmt=fmt, b=neg, wrap=(k == 'sem'))
        _nota(ws, f, nota, desde='C', hasta='C', chars=78.0)
    ws.row_dimensions[S_T['sem']].height = 32
    sem = '$B$%d' % S_T['sem']
    ws.conditional_formatting.add('B%d' % S_T['sem'], FormulaRule(formula=['%s="%s"' % (sem, SEM_PROPIA)], fill=OK))
    ws.conditional_formatting.add('B%d' % S_T['sem'], FormulaRule(formula=['%s="%s"' % (sem, SEM_COMPRA)], fill=AMBAR))
    cel = lambda k: 'B%d' % T[k]
    _dv(ws, 'decimal', [cel(k) for k in ('tp', 'tg', 'mz', 'mm', 'gn')], f1='0', op='greaterThan', msg='Type a number greater than 0.')
    _dv(ws, 'decimal', [cel(k) for k in ('cg', 'tw', 'ti')], f1='0', op='greaterThanOrEqual', msg='Type a number greater than or equal to 0.')
    _dv(ws, 'whole', [cel(k) for k in ('tn', 'ty', 'tpd')], f1='1', f2='1000000', op='between', titulo='Invalid number',
        msg='Type a whole number of 1 or more.')
    _dv(ws, 'whole', [cel('td')], f1='1', f2='366', op='between', titulo='Invalid days', msg='Type a whole number from 1 to 366.')
    _imprimir(ws)


def hoja_escandallo(wb):
    ws = wb.create_sheet(HE)
    for col, w in (('A', 40), ('B', 14), ('C', 7)):
        ws.column_dimensions[col].width = w
    for col in PCOL:
        ws.column_dimensions[col].width = 14
    ws.column_dimensions['N'].width = 44
    _f(ws, 'A1', 'TACO COSTING: 8 TACOS, AN ORDER OF 3 AND A GRINGA AS EXAMPLES, IN OUNCES', b=True, sz=14)
    _f(ws, 'A2', 'Type only in the GREEN cells (prices, ounces, names, tortillas and menu prices). Amounts before sales tax except the menu '
                 'price. Prices are a kit estimate, editable: type your distributor\'s. The meats and the corn tortilla come from the '
                 '"%s" and "%s" sheets.' % (HT, HO), wrap=True)
    ws.merge_cells('A2:N2')
    ws.row_dimensions[2].height = 32
    for col, t in zip('ABD', ('GENERAL DATA', 'Value', 'Notes')):
        _f(ws, '%s4' % col, t, b=True, fill=GRIS)
    for k, rot, fmt, val, nota in ESC_IN:
        f = E[k]
        _f(ws, 'A%d' % f, rot)
        _f(ws, 'B%d' % f, val, fmt=fmt, fill=VERDE, locked=False, wrap=(k == 'modo'))
        _nota(ws, f, nota, desde='D', hasta='N', chars=150.0)
    ws.row_dimensions[E['modo']].height = 32
    for col, t in (('A', 'INGREDIENT (in ounces)'), ('B', '$/lb before tax'), ('C', 'Unit')):
        _f(ws, '%s%d' % (col, E_CAB), t, b=True, fill=GRIS, wrap=True)
    _f(ws, 'N%d' % E_CAB, 'WHERE THE PRICE COMES FROM', b=True, fill=GRIS, wrap=True)
    _f(ws, 'A%d' % E_PVP, 'Menu price ($)', b=True)
    for j, (nom, (pvp, _g, _u)) in enumerate(RECETAS.items()):
        col = PCOL[j]
        _f(ws, '%s%d' % (col, E_CAB), nom, b=True, fill=VERDE, wrap=True, locked=False)
        _f(ws, '%s%d' % (col, E_PVP), pvp, fmt=USD, fill=VERDE, locked=False, b=True)
    ws.row_dimensions[E_CAB].height = 32
    for i, (nom, p, origen) in enumerate(ING_G):
        f = E_G0 + i
        if isinstance(p, str):
            _f(ws, 'A%d' % f, nom)
            if p == 'cono':
                _f(ws, 'B%d' % f, "='%s'!$B$%d" % (HT, S_A['coste_kg']), fmt=USD)
            else:
                _f(ws, 'B%d' % f, "='%s'!$G$%d" % (HT, B0 + int(p[1:])), fmt=USD)
        else:
            _f(ws, 'A%d' % f, nom, fill=VERDE, locked=False)
            _f(ws, 'B%d' % f, p, fmt=USD, fill=VERDE, locked=False)
        _f(ws, 'C%d' % f, 'oz')
        for j, (_n, (_pvp, gr, _ud)) in enumerate(RECETAS.items()):
            _f(ws, '%s%d' % (PCOL[j], f), float(gr.get(nom, 0)), fmt='#,##0.00;-#,##0.00;"-"', fill=VERDE, locked=False)
        _f(ws, 'N%d' % f, origen, i=True)
    _f(ws, 'A%d' % E_UH, 'INGREDIENT (in units)', b=True, fill=GRIS)
    _f(ws, 'B%d' % E_UH, '$/ea before tax', b=True, fill=GRIS)
    _f(ws, 'C%d' % E_UH, 'Unit', b=True, fill=GRIS)
    for i, (nom, p, origen) in enumerate(ING_U):
        f = E_U0 + i
        _f(ws, 'A%d' % f, nom)
        if p is None:
            _f(ws, 'B%d' % f, "=IF($B$%d=\"%s\",'%s'!$B$%d,'%s'!$B$%d)" % (E['modo'], MODOS[0], HO, S_T['c_comp'], HO, S_T['c_prop']), fmt=USD3)
        else:
            _f(ws, 'B%d' % f, p, fmt=USD3, fill=VERDE, locked=False)
        _f(ws, 'C%d' % f, 'ea')
        for j, (_n, (_pvp, _g, ud)) in enumerate(RECETAS.items()):
            _f(ws, '%s%d' % (PCOL[j], f), float(ud.get(nom, 0)), fmt='#,##0.0;-#,##0.0;"-"', fill=VERDE, locked=False)
        _f(ws, 'N%d' % f, origen, i=True)
    _f(ws, 'A%d' % E_R, 'RESULT PER ITEM (ounces and units in the columns above)', b=True, fill=GRIS)
    for j in range(NP):
        _f(ws, '%s%d' % (PCOL[j], E_R), '=%s%d' % (PCOL[j], E_CAB), b=True, fill=GRIS, wrap=True)
    ws.row_dimensions[E_R].height = 32
    oz = _oz_ext()
    for k, rot, fmt in RES:
        f = R[k]
        neg = k in ('c_tot', 'fc')
        _f(ws, 'A%d' % f, rot, b=neg)
        for j in range(NP):
            c = PCOL[j]
            rg = '%s$%d:%s$%d' % (c, E_G0, c, E_G1)
            ru = '%s$%d:%s$%d' % (c, E_U0, c, E_U1)
            ex = {
                'peso': 'SUM(%s)' % rg,
                'c_mat': 'SUMPRODUCT(%s,$B$%d:$B$%d)/%s+SUMPRODUCT(%s,$B$%d:$B$%d)' % (rg, E_G0, E_G1, oz, ru, E_U0, E_U1),
                'c_tot': '%s%d*(1+$B$%d)' % (c, R['c_mat'], E['merm']),
                'pvp_sin': '%s%d/(1+$B$%d)' % (c, E_PVP, E['tax']),
                'fc': '%s%d/%s%d' % (c, R['c_tot'], c, R['pvp_sin']),
                'sem': 'IF(%s%d<=$B$%d,"%s",IF(%s%d<=$B$%d+$B$%d,"%s","%s"))' % (
                    c, R['fc'], E['obj'], FC_OK, c, R['fc'], E['obj'], E['tol'], FC_VIG, FC_ALTO),
                'margen': '%s%d-%s%d' % (c, R['pvp_sin'], c, R['c_tot']),
                'pvp_sug': '%s%d/$B$%d*(1+$B$%d)' % (c, R['c_tot'], E['obj'], E['tax']),
            }[k]
            _f(ws, '%s%d' % (c, f), '=' + ex, fmt=fmt, b=neg, wrap=(k == 'sem'))
    ws.row_dimensions[R['sem']].height = 32
    rng = '%s%d:%s%d' % (PCOL[0], R['sem'], PCOL[-1], R['sem'])
    top = '%s%d' % (PCOL[0], R['sem'])
    for txt, fill in ((FC_OK, OK), (FC_VIG, AMBAR), (FC_ALTO, ROJO)):
        ws.conditional_formatting.add(rng, FormulaRule(formula=['%s="%s"' % (top, txt)], fill=fill))
    # margen por canal de la orden y la gringa
    _f(ws, 'A%d' % E_CH, 'MARGIN BY CHANNEL (order and gringa)', b=True, fill=GRIS)
    for jj, col in CH_COL.items():
        _f(ws, '%s%d' % (col, E_CH), '=%s%d' % (PCOL[jj], E_CAB), b=True, fill=GRIS, wrap=True)
    ws.row_dimensions[E_CH].height = 32
    for k, rot, fmt in CANAL:
        f = CH[k]
        _f(ws, 'A%d' % f, rot, b=k.startswith('plat'))
        for jj, col in CH_COL.items():
            c = PCOL[jj]
            ps = '%s%d' % (c, R['pvp_sin'])
            ex = {
                'sala_e': '%s%d' % (c, R['margen']),
                'sala_p': '%s%d/%s' % (col, CH['sala_e'], ps),
                'llev_e': '%s%d-$B$%d' % (col, CH['sala_e'], E['env']),
                'llev_p': '%s%d/%s' % (col, CH['llev_e'], ps),
                'plat_e': '%s%d-$B$%d*%s' % (col, CH['llev_e'], E['com'], ps),
                'plat_p': '%s%d/%s' % (col, CH['plat_e'], ps),
            }[k]
            _f(ws, '%s%d' % (col, f), '=' + ex, fmt=fmt, b=k.startswith('plat'))
    _f(ws, 'A%d' % E_FIN, 'The food cost in the table is that of the item served dine-in. For takeout the order packaging comes off; on a delivery '
                          'app, the commission over the price before tax comes off too. If the margin on the app falls short, raise the price '
                          'there or sell the order "build-your-own" (tortillas on the side), which travels better. Drinks are not in this costing.',
       i=True, wrap=True)
    ws.merge_cells('A%d:N%d' % (E_FIN, E_FIN))
    ws.row_dimensions[E_FIN].height = 48
    # validaciones
    dcol = lambda cols, fila: ['%s%d' % (c, fila) for c in cols]
    _dv(ws, 'decimal', ['B%d' % E[k] for k in ('tax', 'tol', 'merm', 'com')], f1='0', f2='0.9', op='between', titulo='Invalid percentage',
        msg='Type a percentage between 0 and 0.9 (0.1 = 10 %).')
    _dv(ws, 'decimal', ['B%d' % E['obj']], f1='0.01', f2='0.9', op='between', titulo='Invalid target',
        msg='The target cannot be 0 (the sheet divides by it): type between 0.01 and 0.9.')
    _dv(ws, 'decimal', ['B%d' % E['env']] + ['B%d' % f for i, f in enumerate(range(E_G0, E_G1 + 1)) if not isinstance(ING_G[i][1], str)]
        + ['B%d' % (E_U0 + 1)] + dcol(PCOL, E_PVP)
        + ['%s%d' % (c, f) for c in PCOL for f in list(range(E_G0, E_G1 + 1)) + list(range(E_U0, E_U1 + 1))],
        f1='0', op='greaterThanOrEqual', titulo='Invalid amount', msg='Type a number greater than or equal to 0.')
    _dv(ws, 'list', ['B%d' % E['modo']], f1='"%s"' % ','.join(MODOS), titulo='Invalid option', msg='Pick an option from the list.')
    _imprimir(ws, horizontal=True)


def _instrucciones():
    return [
        ('INSTRUCTIONS: Taco Cost & Trompo Yield Calculator', True),
        ('Three tools in one workbook: the "%s" sheet tells you how many tacos come out of an al pastor cone and what the meat on each one '
         'costs; the "%s" sheet compares buying corn tortillas with making them in house and gives you the break-even in tortillas per day; '
         'the "%s" sheet calculates the cost, food cost, margin and menu price of 8 tacos, an order of 3 and a gringa, with the margin dine-in, '
         'on takeout and on a delivery app. Type only in the GREEN cells; everything else is calculated. Meat prices are in $/lb and '
         'portions in ounces (16 oz per lb, the gray constant).' % (HT, HO, HE), False),
        ('Every default value is a kit estimate, editable: they are examples consistent with the kit\'s business plan case, not anyone\'s '
         'prices. Replace them with your supplier\'s and with what you measure in your kitchen before you decide on prices.', False),
        ('"%s" SHEET' % HT.upper(), True),
        ('1. How to measure the cone. Weigh the raw meat you stack on the trompo (before the marinade) and write down the pounds. At close, '
         'weigh what is left of the cooked cone and note where it goes: reused in a stew or thrown away. With those three numbers and the '
         'tacos you sold, the sheet gives you your real yield.', False),
        ('2. How to adjust the loss. The 35 % cooking and drip loss is a kit estimate with no published source. Measure your first cone: '
         'tacos served × ounces per taco ÷ 16 is the pounds served. Move the loss until the sheet\'s yield matches yours.', False),
        ('3. The most useful results are the cost per served lb and the meat cost per taco, which feed the taco costing. Compare your tacos '
         'per raw lb with the benchmark of a Mexico City stand, about 11 per lb (Milenio). If you serve more ounces, you get fewer tacos.', False),
        ('4. "Raw meat to prepare the N tacos" tells you how much meat to stack for a big order. The other proteins (carnitas, suadero, '
         'cochinita pibil, barbacoa, tinga and campechano) use their own yield after cooking, lower than the cone\'s.', False),
        ('5. Food safety. A stacked, marinated cone is a raw meat preparation: keep it at 41 °F or below until it goes on the broiler, probe it '
         'and log the temperature in your food safety plan. Cook and hold temperatures for al pastor off the trompo are set by your county, '
         'and counties differ: ask your county health department IN WRITING, type the answer in the green "Your county\'s al pastor cook '
         'temp (°F)" cell and keep the reply on file. The kit gives no single figure.', False),
        ('"%s" SHEET' % HO.upper(), True),
        ('1. Type the price of the case, the tortillas per case and the weight of each tortilla, plus the masa harina, machine and tortilla maker '
         'numbers. The tortillas per lb bought are what you need to order from your supplier.', False),
        ('2. The break-even is the number of tortillas a day above which making them pays, counting the person who makes them and the '
         'machine. If you use fewer, buy. With your own tortillas per day, the sheet tells you which suits you.', False),
        ('"%s" SHEET' % HE.upper(), True),
        ('1. Type the sales tax included in your menu prices (0 % if your menu is quoted before tax, the usual US way), the target food cost, '
         'the watch margin, the waste, the packaging and the delivery app commission, and pick whether the corn tortilla is bought or made '
         'with masa harina.', False),
        ('2. In the table, each row is an ingredient with its price per lb before tax (or per unit, for tortillas). The ounces of each item go '
         'in its column; you can change the names, the ounces and the menu prices. The meats take their price from the "%s" sheet. The '
         'conversion from ounces to pounds uses the gray constant (16 oz per lb). The street taco uses two 6 in tortillas.' % HT, False),
        ('3. For each item you see the cost with waste, the price before tax, the food cost with a traffic light, the margin and the menu '
         'price with tax you would need to hit your target. At the bottom, the margin of the order and the gringa dine-in, on takeout and '
         'on a delivery app.', False),
        ('PRICES AND SOURCES', True),
        ('Every default price is an example: "Kit estimate (Oct 2026): replace with your distributor\'s price". The sample menu prices (street '
         'taco $4.00, gringa $11.50) come from the kit\'s business plan case. Sales tax is not part of the menu prices in this sample (0 %); rates '
         'vary by state and city, so ask your accountant. Beer, aguas frescas and other drinks are not in this costing.', False),
        ('Sources: the kit\'s US taco shop research (October 2026); the benchmark of tacos per raw lb comes from Milenio. No figure in this '
         'calculator is published data about your business: replace them with your own.', False),
    ]


def hoja_instrucciones(wb):
    ws = wb.create_sheet(HI)
    ws.column_dimensions['A'].width = 110
    f = 1
    for t, b in _instrucciones():
        _f(ws, 'A%d' % f, t, b=b, sz=13 if f == 1 else 11, wrap=True)
        if not b:
            ws.row_dimensions[f].height = max(16, 15 * math.ceil(len(t) / 120.0))
        f += 2 if b and f > 1 else 1
    f += 1
    for t in ('To edit a cell that is not green: Review → Unprotect Sheet (there is no password).',
              'More templates and kits in the catalog at ' + M.URL_TIENDA, M.FIRMA_EN, VERSION_LINEA, M.AVISO_EN):
        _f(ws, 'A%d' % f, t, wrap=True, i=t == M.AVISO_EN)
        f += 1
    _imprimir(ws)


def docprops(wb):
    pr = wb.properties
    pr.title = '%s · %s' % (TITULO, PRODUCTO)
    pr.subject = '%s · v%s' % (PRODUCTO, VERSION)
    pr.keywords = KEYWORDS
    pr.description = URL
    pr.category = 'AI Chef Pro · Digital products'
    pr.creator = pr.lastModifiedBy = 'AI Chef Pro'


def generar(destino):
    wb = openpyxl.Workbook()
    hoja_trompo(wb)
    hoja_tortilla(wb)
    hoja_escandallo(wb)
    hoja_instrucciones(wb)
    docprops(wb)
    os.makedirs(os.path.dirname(destino), exist_ok=True)
    wb.save(destino)
    import inject_cache as IC                                                      # noqa: E402
    nf, inj, fail = IC.inject(destino)
    print('%s: formulas=%d cache_inyectado=%d fallos_pycel=%d' % (os.path.basename(destino), nf, inj, fail))
    return fail


def valores_expuestos(ruta):
    """Lee de la CACHÉ del xlsx los 5 valores expuestos: {nombre: valor}."""
    wv = openpyxl.load_workbook(ruta, data_only=True)
    return OrderedDict((n, wv[h][c].value) for n, (h, c) in EXPUESTOS.items())


# ---------------------------------------------------------------------------------------------------------------------
# Autotest
# ---------------------------------------------------------------------------------------------------------------------
#: restos de la Pizza EN, euros, español, promesas y reglamentación ES (complementa datos_taq.RX_TAQ_NO si se puede importar)
PROHIBIDOS = [r'pizz', r'neapolit', r'\bdough\b', r'wood[- ]fired', r'margherita', r'mozzarella', r'San Marzano', r'\bVPN\b',
              r'€', r'\bIVA\b', r'APPCC', r'\bkg\b|\bgramos?\b', r'gluten[- ]free', r'authentic', r'mezcal', r'tex-?mex',
              r'\b15-25 ?%', r'\bpico de gallo\b.*\bfree\b', r'\btapas\b|\bcocktails?\b|\bhappy hour\b',
              r'\bfreelance', r'carnet', r'tarifa plana', r'convenio|\bSMI\b|aut[oó]nomo|Veri\*?factu']
CELDA = {}                                                                         # clave de entrada → (hoja, celda)
for _k, _f2 in M_.items():
    CELDA[_k] = (HT, 'B%d' % _f2)
for _i in range(NB):
    for _c, _k in zip('BCDE', ('bk', 'bp', 'br', 'bg')):
        CELDA['%s%d' % (_k, _i)] = (HT, '%s%d' % (_c, B0 + _i))
for _k, _f2 in T.items():
    CELDA[_k] = (HO, 'B%d' % _f2)
for _k, _f2 in E.items():
    CELDA[_k] = (HE, 'B%d' % _f2)
for _i in range(NG):
    CELDA['p%d' % _i] = (HE, 'B%d' % (E_G0 + _i))
CELDA['th'] = (HE, 'B%d' % (E_U0 + 1))
for _j in range(NP):
    CELDA['pvp%d' % _j] = (HE, '%s%d' % (PCOL[_j], E_PVP))
    for _i in range(NG):
        CELDA['g%d_%d' % (_i, _j)] = (HE, '%s%d' % (PCOL[_j], E_G0 + _i))
    for _i in range(NU):
        CELDA['u%d_%d' % (_i, _j)] = (HE, '%s%d' % (PCOL[_j], E_U0 + _i))


def celdas_salida(r):
    """(hoja, celda, valor esperado, etiqueta) de todas las salidas."""
    out = []
    for k in S_A:
        out.append((HT, 'B%d' % S_A[k], r[k], k))
    for i, b in enumerate(r['B']):
        for col, k in zip('FGHI', ('coc', 'ckg', 'tacos', 'ctaco')):
            out.append((HT, '%s%d' % (col, B0 + i), b[k], 'stew%d.%s' % (i, k)))
    for k in S_T:
        out.append((HO, 'B%d' % S_T[k], r['T'][k], 'T.' + k))
    for i in range(NG):
        out.append((HE, 'B%d' % (E_G0 + i), r['pg'][i], 'price%d' % i))
    for i in range(NU):
        out.append((HE, 'B%d' % (E_U0 + i), r['pu'][i], 'unit_price%d' % i))
    for j in range(NP):
        for k in R:
            out.append((HE, '%s%d' % (PCOL[j], R[k]), r['p%d' % j][k], '%s[%d]' % (k, j)))
    for jj, col in CH_COL.items():
        for k in CH:
            out.append((HE, '%s%d' % (col, CH[k]), r['p%d' % jj][k], '%s[%d]' % (k, jj)))
    return out


def _distinto(x, v):
    if isinstance(v, str):
        return x != v
    return not isinstance(x, (int, float)) or abs(float(x) - v) > 0.01


def _checks_logica(nombre, jj, err):
    r = calcular(jj)
    if abs(r['servidos'] - r['rend'] * jj['K']) > 1e-6:
        err.append('%s: servidos ≠ rendimiento × K' % nombre)
    if abs(r['tacos'] * jj['gt'] - r['servidos'] * jj['OZ']) > 0.05:
        err.append('%s: tacos × oz ≠ servidos × 16' % nombre)
    tpd = r['T']['eq']
    if tpd > 0:
        tt = calcular(dict(jj, tpd=tpd))['T']
        if abs(tt['dia_comp'] - tt['dia_prop']) > 0.01:
            err.append('%s: en el equilibrio, coste diario comprada %.4f ≠ propia %.4f' % (nombre, tt['dia_comp'], tt['dia_prop']))
    for j in (J_ORDEN, J_GRINGA):
        p = r['p%d' % j]
        if not (p['plat_e'] < p['llev_e'] < p['sala_e'] and p['plat_p'] < p['llev_p'] < p['sala_p']):
            err.append('%s: margen app < takeout < dine-in no se cumple en la columna %d' % (nombre, j))


def autotest(destino):
    from pycel import ExcelCompiler
    err, n = [], 0
    wv = openpyxl.load_workbook(destino, data_only=True)
    wf = openpyxl.load_workbook(destino)
    base = valores_defecto()
    j2 = dict(base, K=40.0, a=0.10, Pc=3.40, Pa=3.5, pk=3.0, pp=1.0, mc=0.30, mr=0.07, S=2.5, r=0.5, gt=1.25, N=250,
              bk0=24.0, bp1=6.2, br2=0.58, bg3=2.0, tp=52.0, tn=300, tg=1.0, mz=0.5, mm=2.0, cg=0.3, gn=1.1, tw=21000.0, ti=9000.0, ty=5,
              td=280, tpd=900, tax=0.08, obj=0.28, tol=0.03, merm=0.08, modo=MODOS[1], env=0.7, com=0.30, th=0.3, p7=0.9, p14=9.0,
              pvp0=4.25, pvp8=12.0, g0_1=2.0, g3_5=0.3, g1_9=1.0)
    j3 = dict(base, K=15.0, a=0.06, Pc=2.8, mc=0.40, mr=0.04, gt=2.0, N=30, bk3=12.0, br4=0.7, tp=70.0, tn=360, tg=0.8, mz=0.6, mm=1.7,
              cg=0.4, gn=0.8, tw=30000.0, ti=15000.0, ty=8, td=320, tpd=150, obj=0.2, tol=0.1, merm=0.0, env=0.4, com=0.35, th=0.2,
              p9=1.2, pvp9=9.0, g0_8=6.0, g2_2=1.0)
    # 1. caché y pycel del juego por defecto; pycel de los otros dos con set_value
    xl = ExcelCompiler(filename=destino)
    r0 = calcular(base)
    for hoja, c, v, lab in celdas_salida(r0):
        n += 1
        for et, x in (('pycel', xl.evaluate("'%s'!%s" % (hoja, c))), ('caché', wv[hoja][c].value)):
            if _distinto(x, v):
                err.append('defecto · %s!%s (%s): %s %r ≠ esperado %r' % (hoja, c, lab, et, x, v))
    for nombre, jj in (('juego 2', j2), ('juego 3', j3)):
        r = calcular(jj)
        xl2 = ExcelCompiler(filename=destino)
        for hoja, c, v, lab in celdas_salida(r0):
            xl2.evaluate("'%s'!%s" % (hoja, c))
        for k, v in jj.items():
            if v != base[k]:
                h, c = CELDA[k]
                xl2.set_value("'%s'!%s" % (h, c), v)
        for hoja, c, v, lab in celdas_salida(r):
            n += 1
            x = xl2.evaluate("'%s'!%s" % (hoja, c))
            if _distinto(x, v):
                err.append('%s · %s!%s (%s): pycel %r ≠ esperado %r' % (nombre, hoja, c, lab, x, v))
    # 2. identidades de la lógica en los tres juegos (y servidos / tacos también sobre la hoja por defecto)
    for nombre, jj in (('defecto', base), ('juego 2', j2), ('juego 3', j3)):
        n += 3
        _checks_logica(nombre, jj, err)
    oz = xl.evaluate("'%s'!B%d" % (HT, M_['OZ']))
    k_ = xl.evaluate("'%s'!B%d" % (HT, M_['K']))
    sv, rd = xl.evaluate("'%s'!B%d" % (HT, S_A['servidos'])), xl.evaluate("'%s'!B%d" % (HT, S_A['rend']))
    tc, gt = xl.evaluate("'%s'!B%d" % (HT, S_A['tacos'])), xl.evaluate("'%s'!B%d" % (HT, M_['gt']))
    n += 2
    if abs(sv - rd * k_) > 1e-6 or abs(tc * gt - sv * oz) > 0.05:
        err.append('hoja: servidos ≠ rendimiento × K o tacos × oz ≠ servidos × 16')
    # 3. valores con sentido
    pastor = r0['p%d' % J_PASTOR]
    if not 0.20 <= pastor['fc'] <= 0.30:
        err.append('food cost del taco al pastor %.1f %% fuera de 20-30 %%' % (pastor['fc'] * 100))
    if abs(r0['rend'] - 0.648) > 0.001 or not 206 <= r0['tacos'] <= 209 or abs(r0['tacos_lb'] - 6.9) > 0.1:
        err.append('rendimiento %.4f / tacos %.1f / tacos por lb %.2f fuera de lo esperado (64,8 %% · 207 · 6,9)'
                   % (r0['rend'], r0['tacos'], r0['tacos_lb']))
    if abs(r0['coste_kg'] - 5.00) > 0.01 or abs(r0['carne_taco'] - 0.469) > 0.005:
        err.append('coste/lb servida %.3f o carne por taco %.4f fuera de 5,00 / 0,47' % (r0['coste_kg'], r0['carne_taco']))
    if not 0.95 <= pastor['c_tot'] <= 1.05:
        err.append('coste del taco al pastor %.4f fuera de 0,95-1,05' % pastor['c_tot'])
    if not 640 <= r0['T']['eq'] <= 680 or r0['T']['sem'] != SEM_COMPRA:
        err.append('equilibrio %.0f fuera de 640-680 o semáforo ≠ compra con 650 al día' % r0['T']['eq'])
    if calcular(dict(base, tpd=900))['T']['sem'] != SEM_PROPIA:
        err.append('el semáforo de la tortilla no pasa a «Make them» por encima del equilibrio')
    for j in range(NP):
        if not 0.10 <= r0['p%d' % j]['fc'] <= 0.40:
            err.append('food cost de %s fuera de 10-40 %%: %.1f' % (list(RECETAS)[j], r0['p%d' % j]['fc'] * 100))
    # 4. los 5 valores expuestos: caché frente a calcular()
    esperados = {'TACOS_POR_CONO': r0['tacos'], 'RENDIMIENTO_TROMPO': r0['rend'], 'COSTE_TACO_PASTOR': pastor['c_tot'],
                 'FC_TACO_PASTOR': pastor['fc'], 'EQUILIBRIO_TORTILLAS_DIA': r0['T']['eq']}
    expuestos = valores_expuestos(destino)
    for nombre, e_ in esperados.items():
        n += 1
        if _distinto(expuestos[nombre], e_):
            err.append('expuesto %s: caché %r ≠ %r' % (nombre, expuestos[nombre], e_))
    # 5. formato, protección, DV, fórmulas, textos
    if motor.gate_formatos(wf, FICHERO):
        err.append('gate_formatos: %r' % motor.gate_formatos(wf, FICHERO)[:3])
    dvh = dict((w.title, dict((str(rng), dv) for dv in w.data_validations.dataValidation for rng in dv.sqref.ranges))
               for w in (wf[HT], wf[HO], wf[HE]))
    for hoja, k in ((HT, 'K'), (HT, 'N'), (HT, 'mc'), (HT, 'ct'), (HO, 'tpd'), (HO, 'td'), (HO, 'tp'), (HE, 'obj'), (HE, 'modo'), (HE, 'tax')):
        c = CELDA[k][1]
        if c not in dvh[hoja]:
            err.append('DV ausente en %s!%s' % (hoja, c))
    dv = dvh[HT].get(CELDA['N'][1])
    if dv is None or dv.type != 'whole':
        err.append('DV del número de tacos (entero) ausente o incorrecta')
    dv = dvh[HE].get(CELDA['modo'][1])
    if dv is None or dv.type != 'list':
        err.append('DV de lista del modo de tortilla ausente')
    ct = wf[HT][CELDA['ct'][1]]
    if ct.value is not None or ct.protection.locked or ct.fill.fgColor.rgb not in ('00E8F5E9', 'FFE8F5E9') \
            or "cook temp (°F)" not in str(wf[HT]['A%d' % M_['ct']].value):
        err.append('campo «cook temp» del condado: debe existir, vacío, verde y desbloqueado')
    textos = []
    for w in wf.worksheets:
        if w.page_setup.paperSize != 1 or not w.sheet_properties.pageSetUpPr.fitToPage or not w.protection.sheet \
                or w.oddFooter.center.text != FOOTER:
            err.append('%s: Letter + ajuste + protección + pie' % w.title)
        for row in w.iter_rows():
            for c in row:
                v = c.value
                if v == '':
                    err.append('%s!%s: cadena vacía' % (w.title, c.coordinate))
                if isinstance(v, str) and c.data_type != 'f':
                    textos.append((w.title, c.coordinate, v))
                if isinstance(v, str) and c.data_type == 'f':
                    textos += [(w.title, c.coordinate + ' (fórmula)', lit) for lit in re.findall(r'"((?:[^"]|"")*)"', v)
                               if any(ch.isalpha() for ch in lit)]
                if c.data_type == 'f':
                    if wv[w.title][c.coordinate].value is None:
                        err.append('%s!%s: fórmula sin caché' % (w.title, c.coordinate))
                    sin_refs = re.sub(r'"[^"]*"', '', re.sub(r"'[^']+'!|\$?[A-Z]{1,2}\$?\d+", '', str(c.value)))
                    for num in re.findall(r'(?<![A-Za-z])\d+\.?\d*', sin_refs):
                        if num not in ('0', '1'):
                            err.append('%s!%s: número tecleado %s en %s' % (w.title, c.coordinate, num, c.value))
                if c.fill.fgColor.rgb in ('00E8F5E9', 'FFE8F5E9') and c.protection.locked:
                    err.append('%s!%s: verde bloqueada' % (w.title, c.coordinate))
                if c.fill.fgColor.rgb in ('00E8F5E9', 'FFE8F5E9') and c.data_type == 'f':
                    err.append('%s!%s: verde con fórmula' % (w.title, c.coordinate))
                fmt = c.number_format or ''
                if '€' in fmt or '$' in fmt:
                    err.append('%s!%s: símbolo de moneda en el formato (D8)' % (w.title, c.coordinate))
        for dv_ in w.data_validations.dataValidation:
            for t in (dv_.error, dv_.errorTitle, dv_.prompt, dv_.promptTitle):
                if t:
                    textos.append((w.title, 'DV %s' % dv_.sqref, t))
    p = wf.properties
    for campo in ('title', 'subject', 'keywords', 'description', 'category'):
        textos.append(('docProps', campo, getattr(p, campo) or ''))
    try:
        import datos_taq as D
        rx_taq = [x[0] for x in D.RX_TAQ_NO]
    except Exception as e:                                                         # noqa: BLE001
        rx_taq = []
        print('AVISO: no se pudo importar datos_taq.RX_TAQ_NO (%s)' % e)
    n_es = n_pz = 0
    for h, c, t in textos:
        r = [x for x in M.restos_espanol(t) if not x[0].startswith('oficio')]
        if r:
            n_es += 1
            err.append('%s!%s: restos de español %s' % (h, c, r[:1]))
        if M.no_latinos(t):
            err.append('%s!%s: caracteres no latinos' % (h, c))
        for rx in PROHIBIDOS + rx_taq:
            if re.search(rx, t, re.I if not rx.startswith('(?i)') else 0):
                n_pz += 1
                err.append('%s!%s: prohibido /%s/ en %r' % (h, c, rx, t[:60]))
        if re.search(r'(?i)\bvalue\b.*\$\d|\bsave \$|total value|\bwas \$|\d+\s?% off\b', t):
            err.append('%s!%s: ancla de precio' % (h, c))
    todo = '\n'.join(t for _h, _c, t in textos)
    for frase in ('16 oz per lb', 'Sales tax included in your menu prices (%)', "Your county's al pastor cook temp (°F)",
                  'IN WRITING', SEM_COMPRA, SEM_PROPIA, FC_OK, FC_VIG, FC_ALTO, 'Kit estimate',
                  "Kit estimate (Oct 2026): replace with your distributor's price"):
        if frase not in todo:
            err.append('falta el texto %r' % frase)
    if M.FIRMA_EN not in todo or VERSION_LINEA not in todo or M.AVISO_EN not in todo:
        err.append('falta la firma, la línea de versión o el aviso')
    if wf[HE]['A5'].value != 'Sales tax included in your menu prices (%)' or wf[HE]['B5'].value != 0:
        err.append('Taco Costing!B5 no es el sales tax = 0 %')
    if wf.sheetnames != [HT, HO, HE, HI]:
        err.append('pestañas %r' % wf.sheetnames)
    if NG + NU < 16:
        err.append('menos de 16 ingredientes')
    if p.creator != 'AI Chef Pro' or p.description != URL or 'AI Chef Pro' not in (p.keywords or '') \
            or not (p.title or '').startswith(TITULO):
        err.append('docProps: creator=%r title=%r description=%r keywords=%r' % (p.creator, p.title, p.description, p.keywords))
    for x in err[:80]:
        print('ERROR', x)
    print('autotest taco cost & trompo yield calculator: %d comprobaciones (pycel + caché por defecto; pycel en 2 juegos más; identidades) '
          'sobre %d textos · %s' % (n, len(textos), 'VERDE' if not err else 'ROJO (%d errores)' % len(err)))
    print('  cone: %.0f lb → %.2f lb served · yield %.1f %% · %.1f tacos (%.2f per raw lb) · $%.3f per served lb · meat $%.3f per taco'
          % (base['K'], r0['servidos'], r0['rend'] * 100, r0['tacos'], r0['tacos_lb'], r0['coste_kg'], r0['carne_taco']))
    for nombre, (h, c) in EXPUESTOS.items():
        print('  %-26s %s!%-5s %s' % (nombre, h, c, expuestos[nombre]))
    for j, nom in enumerate(RECETAS):
        q = r0['p%d' % j]
        print('  %-28s cost $%.3f · food cost %5.1f %% · %s' % (nom, q['c_tot'], q['fc'] * 100, q['sem']))
    for j in (J_ORDEN, J_GRINGA):
        q = r0['p%d' % j]
        print('  channel %-24s dine-in $%.2f (%.0f %%) · takeout $%.2f (%.0f %%) · app $%.2f (%.0f %%)'
              % (list(RECETAS)[j], q['sala_e'], q['sala_p'] * 100, q['llev_e'], q['llev_p'] * 100, q['plat_e'], q['plat_p'] * 100))
    return 1 if err else 0


if __name__ == '__main__':
    ap = argparse.ArgumentParser()
    ap.add_argument('--salida', required=True, help='output folder')
    ap.add_argument('--autotest', action='store_true')
    a = ap.parse_args()
    destino = os.path.join(a.salida, FICHERO)
    if not a.autotest and generar(destino):
        sys.exit('ERROR: fallos de pycel en inject_cache')
    sys.exit(autotest(destino))
