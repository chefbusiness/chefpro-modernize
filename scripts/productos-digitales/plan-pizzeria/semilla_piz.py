#!/usr/bin/env python3
"""
semilla_piz.py — Plan de Negocio: Pizzería Napolitana (ES 1.0; SPEC D4/D13/D14): prepara el ORIGEN que
`planes-v2_0/main.py` convierte en el plan financiero y el checklist de la pizzería. Copia adaptada de
`plan-heladeria/semilla_hel.py` (F1-camino §3).

Copia los dos xlsx del Tapas Bar ES 2.2 publicados (`dl/plan-negocio-tapas-bar/`, SOLO LECTURA; molde A-β + checklist
C2) con los nombres de fichero de la pizzería y escribe lo que el módulo de contenido no alcanza (F1-camino §1):
  · plan: valores de `0. Supuestos` (grupo_a conserva lo verde que trae el libro), títulos de hoja (Inversión A1/A2/A3,
    PyG A2, Punto Equilibrio A2), las 18 partidas de inversión (filas 7-24: rótulo, importe, nota, ¿IVA?), los 9
    costes fijos propios del P&L (filas 26-34: rótulo e importe; la nota la pone el contenido), `PyG!G13` por VALOR
    (era fórmula en el tapas bar) y los docProps description/keywords;
  · checklist: título, pestañas, cabeceras de fase y las tareas de `datos_checklist.py` con el pie «Tareas
    completadas:» debajo; `Instrucciones!A3` con el pid; DV, formato condicional y contadores los rehace el motor.
La semilla del tapas bar no trae línea partida en `PyG!H`: no hay nada que vaciar (el motor limpia la columna H).
Corre en el VPS (openpyxl). Después: `main.py --producto plan-negocio-pizzeria --dry-run --origen <salida>`.

    /root/venv-guias/bin/python semilla_piz.py --salida /root/tmp-piz-f2/motor/origen

Sesión Claude Code, 25-oct-2026.
"""
import argparse
import copy
import importlib
import os
import shutil
import sys

import openpyxl

import datos_checklist as DC

AQUI = os.path.dirname(os.path.abspath(__file__))
REPO = os.path.abspath(os.path.join(AQUI, '..', '..', '..'))
TAP = os.path.join(REPO, 'astro-site', 'public', 'dl', 'plan-negocio-tapas-bar')
PLAN_TAP, CHECK_TAP = 'plan-financiero-tapas-bar.xlsx', 'checklist-apertura-tapas-bar.xlsx'
PLAN, CHECK = 'plan-financiero-pizzeria.xlsx', 'checklist-apertura-pizzeria.xlsx'
PID = 'plan-negocio-pizzeria'

TITULOS_PLAN = {
    ('Inversión Inicial', 'A1'): 'PLAN DE NEGOCIO: PIZZERÍA NAPOLITANA',
    ('Inversión Inicial', 'A2'): 'Inversión inicial estimada — pizzería con horno de leña, 50 plazas y terraza',
    ('Inversión Inicial', 'A3'): 'Ajusta las cifras a tu proyecto: son las de un caso tipo en una ciudad media, sin '
                                 'traspaso. Pon tus presupuestos reales en las celdas de importe.',
    ('PyG 3 Años', 'A2'): 'Escenario realista — pizzería con horno de leña, 50 plazas y terraza',
    ('Punto Equilibrio', 'A2'): 'Punto de equilibrio mensual y diario — pizzería con horno de leña',
}

#: Filas 7-24 de 'Inversión Inicial' (molde del tapas bar): (rótulo, importe, nota, ¿lleva IVA?). Research §3: 19
#: partidas en 18 filas (la 6 «Rótulo y fachada», 2.500 €, se funde con la 2 «Obra»: SPEC D8), 124.200 € (obra 60.500 ·
#: maquinaria 50.700 · no amortizable 13.000) + imprevistos del 10 % sobre la obra (fila 27, por fórmula) = 130.250 €.
#: Sin palabras de recuento en el rótulo (motor.RX_RECUENTO), sin paréntesis con número y sin «tapa». ⚠️ «campana»
#: a secas la acentúa `motor.TILDES` a «campaña»: siempre «campana extractora» (motor.EXCEPCIONES_CAMPANA).
INVERSION = [
    ('Proyecto técnico de actividad y de salida de humos, y dirección de obra', 4000,
     'Proyecto de la actividad y de la chimenea, memoria para la licencia o la declaración responsable y certificado '
     'final de obra', 'Sí'),
    ('Obra y acondicionamiento del local, rótulo y fachada', 32500,
     'Obrador alicatado, suelo antideslizante, aseos (uno adaptado), falsos techos y aislamiento acústico, más el '
     'rótulo y la fachada (unos 2.500 €). Estimación del kit: pide dos presupuestos', 'Sí'),
    ('Chimenea de doble pared hasta cubierta, filtro de partículas y bancada del horno', 9000,
     'Conducto exclusivo hasta la cubierta y, con leña, filtro de partículas: Madrid pide superar en 1 m lo que haya '
     'en 15 m (Ordenanza 4/2021). Ningún fabricante da precio por metro: pide oferta con tu proyecto', 'Sí'),
    ('Instalación eléctrica trifásica, cuadro y boletín', 6000,
     'Unos 25 kW en trifásica para el frío, la amasadora, el lavavajillas, la extracción y la climatización. Pregunta '
     'la potencia disponible ANTES de firmar: una acometida nueva tarda', 'Sí'),
    ('Fontanería, separador de grasas, climatización y extracción', 9000,
     'Agua fría y caliente en el obrador, desagües con separador de grasas, climatización de la sala y extracción de '
     'la cocina de apoyo', 'Sí'),
    ('Horno de leña de cúpula baja de 120-130 cm de diámetro', 9000,
     'Para unas seis pizzas a la vez. Los prefabricados profesionales con precio publicado van de 2.750 a 5.505 € '
     '(octubre de 2026); los artesanos napolitanos no publican tarifa: pide oferta con el montaje', 'Sí'),
    ('Amasadora de espiral profesional', 3500,
     'Para 30-40 kg de masa: el sábado se cubre en dos amasadas. Las de entrada empiezan en unos 860 € sin IVA; '
     'una de gama profesional cuesta bastante más', 'Sí'),
    ('Mesa fría de pizza con encimera de granito y vitrina', 2200,
     'Mesa refrigerada de dos o tres puertas con vitrina de ingredientes: unos 1.950 € sin IVA en la tarifa de un '
     'distribuidor (octubre de 2026)', 'Sí'),
    ('Armarios refrigerados para las cajas de bolas y arcón congelador', 5000,
     'Armarios ventilados para las cajas de bolas y los ingredientes (unos 900 € sin IVA el de 400 l) y un arcón',
     'Sí'),
    ('Campana extractora con motor, freidora y fogón de apoyo', 3000,
     'Campana extractora de 1,5-2 m con motor (unos 815 € sin IVA) para la cocina de entrantes, frituras y salsas', 'Sí'),
    ('Lavavajillas de capota, fregaderos, lavamanos, mesas inox y estanterías', 6000,
     'Lavavajillas, fregaderos, lavamanos de accionamiento no manual, mesas de acero inoxidable y estanterías', 'Sí'),
    ('Utillaje y menaje: palas, cajas de bolas, básculas, vajilla y uniformes', 5000,
     'Palas de entrada y de giro (unos 70-80 € cada una), cajas de bolas con cierre, básculas, termómetro de '
     'infrarrojos, vajilla, uniformes y la cartelería de alérgenos', 'Sí'),
    ('Barra: cafetera y molino, botellero frío y equipo de cerveza', 6000,
     'Si la cervecera cede el equipo de tiraje, esta partida baja', 'Sí'),
    ('Mobiliario de sala y terraza', 9000,
     'Mesas y sillas para 50 plazas en sala y 24 en la terraza de temporada', 'Sí'),
    ('TPV con Veri*factu, datáfonos, impresora de comandas y tablet', 2000,
     'Software de facturación conforme a Veri*factu, datáfonos, impresora de cocina y la tablet de las plataformas',
     'Sí'),
    ('Constitución de la SL, licencia, tasas, registro sanitario y APPCC', 3500,
     'SL unipersonal, licencia o declaración responsable, tasas municipales, registro sanitario autonómico y '
     'redacción del plan APPCC. Las tasas no llevan IVA; los honorarios sí: cambia la columna de al lado si tu '
     'presupuesto es sobre todo de honorarios', 'No'),
    ('Marca, web, fotos, formación de pizzaiolo y lanzamiento', 5500,
     'Identidad, web con la carta, sesión de fotos, curso de pizzaiolo para el segundo pizzaiolo y campaña de '
     'apertura', 'Sí'),
    ('Stock inicial de harina, tomate, quesos, bebida, leña y cajas', 4000,
     'Harina, tomate, quesos y embutidos, la bodega y la cerveza, la primera leña seca y las cajas de cartón', 'Sí'),
]
assert len(INVERSION) == 18 and sum(x[1] for x in INVERSION) == 124200, sum(x[1] for x in INVERSION)

#: Filas 26-34 de 'PyG 3 Años' (costes fijos propios). SPEC D8: 9 filas, 17.600 €/año, una por fila del molde
#: (no se suprimen ni se añaden filas: desplazaría el P&L), + 1.200 € de tasa de terraza en la fila de los residuos
#: (calibrado D9, vuelta 2: demo §2.11.5) = 18.800 €. Sin palabras de recuento en el rótulo. Las claves de
#: `contenido….FIJOS` son estos rótulos normalizados.
FIJOS = {
    26: ('Gestoría laboral, fiscal y contable', 2400),
    27: ('Mantenimiento del horno, la campana extractora y el frío', 2400),
    28: ('Marketing y publicidad', 4000),
    29: ('Software, TPV, teléfono e internet', 1500),
    30: ('Limpieza, lavandería, EPI, DDD y análisis de autocontrol', 2000),
    31: ('Leña para el horno', 3600),
    32: ('Tasa de terraza y gestión de residuos (pon las de tu ayuntamiento)', 1800),    # calibrado D9, vuelta 2
    33: ('Derechos de autor por música ambiental (SGAE y AGEDI-AIE)', 600),
    34: ('Prevención de riesgos laborales y vigilancia de la salud', 500),
}
assert sum(x[1] for x in FIJOS.values()) == 18800


def _contenido():
    sys.path.insert(0, os.path.join(os.path.dirname(AQUI), 'planes-v2_0'))
    import grupo_a
    return grupo_a, importlib.import_module('contenido_plan_negocio_pizzeria.a')


def _supuestos(wb):
    """Los valores de `0. Supuestos` del caso pizzería: `grupo_a` CONSERVA lo que el libro ya trae en las celdas
    verdes y sólo usa el valor del contenido en una celda vacía; sin este paso el caso seguiría siendo el del
    tapas bar (56 × 18 € × 300)."""
    grupo_a, contenido = _contenido()
    coords = dict((t[0], t[1]) for t in grupo_a.SUPUESTOS_BASE)
    ws = wb['0. Supuestos']
    escritas = 0
    for clave, t in contenido.SUPUESTOS.items():
        if t[2] is None:
            continue
        assert clave in coords, clave
        ws[coords[clave]] = t[2]
        escritas += 1
    return escritas


def _docprops(wb):
    """El motor no reescribe `description` ni `keywords` y heredaba los del tapas bar."""
    wb.properties.description = 'aichef.pro/' + PID
    wb.properties.keywords = 'plan de negocio pizzería napolitana, horno de leña, AI Chef Pro'


def plan(salida):
    destino = os.path.join(salida, PLAN)
    shutil.copy2(os.path.join(TAP, PLAN_TAP), destino)
    wb = openpyxl.load_workbook(destino)
    print('  supuestos escritos:', _supuestos(wb))
    _docprops(wb)
    for (hoja, coord), texto in TITULOS_PLAN.items():
        wb[hoja][coord] = texto
    ws = wb['Inversión Inicial']
    for i, (rot, imp, nota, iva) in enumerate(INVERSION):
        r = 7 + i
        assert not str(ws.cell(r, 1).value).lower().startswith(('fianza', 'alquiler', 'imprevistos')), r
        ws.cell(r, 1).value, ws.cell(r, 2).value, ws.cell(r, 4).value, ws.cell(r, 5).value = rot, imp, nota, iva
    assert str(ws['A25'].value).lower().startswith('fianza'), ws['A25'].value
    assert str(ws['A26'].value).lower().startswith('alquiler'), ws['A26'].value
    assert str(ws['A27'].value).lower().startswith('imprevistos'), ws['A27'].value
    ws = wb['PyG 3 Años']
    assert str(ws['A25'].value).startswith('Seguros') and str(ws['A35'].value).startswith('Amortización'), (
        ws['A25'].value, ws['A35'].value)
    for r, (rot, imp) in FIJOS.items():
        ws.cell(r, 1).value, ws.cell(r, 2).value = rot, imp
        ws.cell(r, 6).value = None          # la nota la pone el contenido (FIJOS)
    # IVA soportado de la compra de comida: en el tapas bar era una FÓRMULA (= IVA de restauración); aquí es la celda
    # verde de `IVA_COMPRAS` (6,4 %). Se escribe el valor en el origen para que el motor no la cuente como una
    # fórmula perdida (control «el post-proceso PERDIÓ fórmulas que el fichero ya tenía»; gotcha de heladería y finca).
    _g, contenido = _contenido()
    assert str(ws['A13'].value).startswith('Coste de mercancía — comida'), ws['A13'].value
    assert str(ws['G13'].value).startswith('='), ws['G13'].value
    ws['G13'] = contenido.IVA_COMPRAS['comida']
    wb.save(destino)
    return destino


def _copiar_estilo(origen, destino):
    destino.font, destino.border = copy.copy(origen.font), copy.copy(origen.border)
    destino.fill, destino.alignment = copy.copy(origen.fill), copy.copy(origen.alignment)
    destino.number_format, destino.protection = origen.number_format, copy.copy(origen.protection)


def checklist(salida):
    destino = os.path.join(salida, CHECK)
    shutil.copy2(os.path.join(TAP, CHECK_TAP), destino)
    wb = openpyxl.load_workbook(destino)
    fases = [ws for ws in wb.worksheets if ws.title[:2] in ('F1', 'F2', 'F3', 'F4', 'F5', 'F6')]
    assert len(fases) == 6
    for ws, (titulo, cabecera, tareas) in zip(fases, DC.FASES):
        ws.title = titulo
        ws['A1'], ws['A2'] = DC.TITULO, cabecera
        assert ws['B4'].value and 'TRÁMITE' in ws['B4'].value.upper(), ws['B4'].value
        modelo = [copy.copy(ws.cell(5, c)) for c in range(1, 6)]
        alto = ws.row_dimensions[5].height
        pie = [copy.copy(ws.cell(r, c)) for r in range(5, ws.max_row + 1) for c in range(1, 6)
               if isinstance(ws.cell(r, c).value, str) and 'completad' in ws.cell(r, c).value.lower()]
        assert pie, ws.title
        fila_pie = pie[0].row
        estilo_pie = [copy.copy(ws.cell(fila_pie, c)) for c in range(1, 6)]
        # vaciar la zona de tareas y el pie (valores, DV y formato condicional los rehace el motor)
        for r in range(5, ws.max_row + 1):
            for c in range(1, 6):
                ws.cell(r, c).value = None
            ws.row_dimensions[r].height = None
        ws.data_validations.dataValidation = []
        ws.conditional_formatting = type(ws.conditional_formatting)()
        for i, (tarea, resp, plazo, nota) in enumerate(tareas):
            r = 5 + i
            for c, v in enumerate(('☐', tarea, resp, plazo, nota), start=1):
                cel = ws.cell(r, c)
                _copiar_estilo(modelo[c - 1], cel)
                cel.value = v
            ws.row_dimensions[r].height = alto
        r = 5 + len(tareas) + 1
        for c in range(1, 6):
            _copiar_estilo(estilo_pie[c - 1], ws.cell(r, c))
        ws.cell(r, 2).value = 'Tareas completadas:'
        ws.cell(r, 3).value = '=COUNTIF(A5:A%d,"✓")' % (4 + len(tareas))
        ws.cell(r, 4).value = 'de'
        ws.cell(r, 5).value = '=COUNTIF(B5:B%d,"?*")-COUNTIF(A5:A%d,"N/A")' % (4 + len(tareas), 4 + len(tareas))
    # `motor.asegurar_instrucciones` no reescribe la hoja si ya existe: la del tapas bar nombraba su producto
    ins = wb['Instrucciones']
    assert 'plan-negocio-tapas-bar' in str(ins['A3'].value), ins['A3'].value
    ins['A3'] = 'Este fichero forma parte del «Plan de Negocio: Pizzería Napolitana» de AI Chef Pro.'   # m13: el nombre, no el slug
    _docprops(wb)
    wb.save(destino)
    return destino


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--salida', required=True)
    a = ap.parse_args()
    if os.path.isdir(a.salida):
        shutil.rmtree(a.salida)
    os.makedirs(a.salida)
    print('plan      →', plan(a.salida))
    print('checklist →', checklist(a.salida), '(%d tareas: %s)' % (
        sum(len(t) for _, _, t in DC.FASES), '/'.join(str(len(t)) for _, _, t in DC.FASES)))


if __name__ == '__main__':
    main()
