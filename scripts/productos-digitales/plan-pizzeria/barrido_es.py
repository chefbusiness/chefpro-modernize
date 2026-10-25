#!/usr/bin/env python3
"""barrido_es.py — Pizzería Napolitana ES 1.0 (SPEC D12; copia de `plan-heladeria/barrido_es.py`): barre los xlsx/docx que
se le pasen buscando los prohibidos de la SPEC (España 2026, validado, freelance, carné de manipulador, tarifa plana,
CNAE, RGSEAA para el minorista, licencia de apertura, RD 895/2013, RD 393/2007, neto «15-25 %», versión anterior…),
los restos de los moldes (Food Truck, Tapas Bar —la semilla—, Heladería —el formato—) y las frases del motor sin
traducir («alcohol que sale por delivery», «pan común»): la lista es `tandas_docx.PROHIBIDOS` (F1-camino §3), más
«Campaña» por «campana» (`RX_CAMPANA`) y «pública» por el verbo «publica» (`RX_PUBLICA`). SIN los
restos de panadería del gemelo: harina, masa, levadura y amasado son el oficio. Lista celda/párrafo y texto; en las
fórmulas sólo mira los literales entre comillas. Corre en el VPS. Uso: barrido_es.py <fichero>… · exit 1 si hay
coincidencias. Sesión Claude Code, 25-oct-2026."""
import os
import re
import sys

import docx
import openpyxl

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import tandas_docx as T                                                            # noqa: E402

#: «Campaña» por «campana» (motor.TILDES acentúa «campana» suelta: medido en el dry-run del 25-oct, A16 y PyG!A27, que
#: además dejó la partida FUERA de la amortización). En este libro nunca hay «campaña» legítima salvo la de apertura,
#: lanzamiento, marketing, publicidad o Navidad: excepción explícita (aviso del orquestador, 25-oct).
RX_CAMPANA = r'[Cc]ampa[ñÑ]a(?! de (apertura|lanzamiento|marketing|publicidad|Navidad))'
#: «pública» por «publica» (la misma trampa de motor.TILDES con el verbo: «Nadie pública precio», Inversión!D9, revisión
#: final m1, 25-oct): con tilde tras un sujeto o un pronombre átono, o delante de «precio» / «tarifa», es el verbo mal
#: acentuado. Sin «lo» ni «la» («lo público», «la pública» son legítimos).
RX_PUBLICA = r'\b(nadie|se|le|les|te|nos|quien) públic[ao]n?\b|\bpúblic[ao] (precios?|tarifas?)\b'
RX = re.compile('|'.join(T.PROHIBIDOS + [RX_CAMPANA, RX_PUBLICA]), re.I)
n = 0
for p in [x for x in sys.argv[1:] if not x.startswith('--')]:
    if p.endswith('.xlsx'):
        wb = openpyxl.load_workbook(p)
        it = [(ws.title + '!' + c.coordinate, c.value) for ws in wb.worksheets for row in ws.iter_rows() for c in row
              if isinstance(c.value, str)]
        it += [('pestaña', ws.title) for ws in wb.worksheets]
        pr = wb.properties
        it += [('docProps.' + k, getattr(pr, k)) for k in ('title', 'subject', 'description', 'keywords', 'creator')
               if isinstance(getattr(pr, k), str)]
    else:
        d = docx.Document(p)
        it = [('p%d' % i, x.text) for i, x in enumerate(d.paragraphs)]
        it += [('tabla%d' % i, c.text) for i, t in enumerate(d.tables) for row in t.rows for c in row.cells]
        cp = d.core_properties
        it += [('docProps.' + k, getattr(cp, k)) for k in ('title', 'subject', 'comments', 'keywords', 'author')
               if isinstance(getattr(cp, k), str)]
    for loc, t in it:
        txt = ' '.join(re.findall(r'"([^"]*)"', t)) if t.startswith('=') else t
        m = RX.search(txt)
        if m:
            n += 1
            print('%s · %s · «%s» · …%s…' % (p.split('/')[-1][:28], loc, m.group(0), txt[max(0, m.start() - 50):m.end() + 50]))
print('barrido: %d coincidencias' % n)
sys.exit(1 if n else 0)
