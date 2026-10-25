#!/usr/bin/env python3
"""
calculadora_piz_es.py — Plan de Negocio: Pizzería Napolitana ES 1.0 (SPEC D11 y D15): `calculadora-masa-escandallo-pizzeria.xlsx`.
Hojas `Masa` · `Escandallo` · `Instrucciones`. Molde de código: `plan-finca-eventos/calculadora_fin_es.py` (helpers, inject_cache,
autotest con pycel, protección sin contraseña, verdes desbloqueadas, A4, docProps AI Chef Pro).
  · Masa, bloque A: por LITRO DE AGUA con los gramajes del pliego de la ETG (Reg. 97/2010, anexo II): nº de bolas, peso de bola,
    merma, harina / sal / levadura en g por litro, €/kg → litros, agua, harina, sal, levadura, % panadero, coste y coste por bola,
    y un semáforo «dentro de los gramajes del pliego» (que NO convierte la pizza en «pizza napolitana»: D11).
  · Masa, bloque B: por % panadero (hidratación, sal y levadura sobre la harina) → gramos, coste de la masa y por bola.
  · Escandallo: 22 ingredientes × 8 pizzas en gramos → coste, PVP sin IVA, food cost con semáforo, margen, PVP recomendado para
    el objetivo y food cost para llevar con la caja. La carta de ejemplo no usa el nombre «napolitana».
Corre en el VPS (openpyxl + pycel):

    /root/venv-guias/bin/python calculadora_piz_es.py --salida <dir>      # genera + inject_cache + autotest
    /root/venv-guias/bin/python calculadora_piz_es.py --salida <dir> --autotest

Sesión Claude Code, 25-oct-2026.
"""
import argparse
import copy
import math
import os
import re
import sys
from collections import OrderedDict

import openpyxl
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill, Protection
from openpyxl.worksheet.datavalidation import DataValidation

AQUI = os.path.dirname(os.path.abspath(__file__))
PD = os.path.dirname(AQUI)
sys.path.insert(0, AQUI)
import piz_version                                                                 # noqa: E402,F401  (parchea motor: 1.0 · octubre 2026)
import motor                                                                       # noqa: E402

PID = 'plan-negocio-pizzeria'
FICHERO = 'calculadora-masa-escandallo-pizzeria.xlsx'
TITULO = 'Calculadora de masa y escandallo por pizza'
HM, HE, HI = 'Masa', 'Escandallo', 'Instrucciones'
VERDE = PatternFill('solid', fgColor='E8F5E9')
GRIS = PatternFill('solid', fgColor='F2F2F2')
ROJO = PatternFill('solid', bgColor='F8CBAD', fgColor='F8CBAD')
AMBAR = PatternFill('solid', bgColor='FFE699', fgColor='FFE699')
OK = PatternFill('solid', bgColor='C6E0B4', fgColor='C6E0B4')
FUENTE = 'Calibri'
EUR, ENT, PCT, DEC, GR, EUR3 = '#,##0.00 €', '#,##0', '0.0%', '#,##0.0', '#,##0.0', '#,##0.000 €'
METODOS = ['Pliego (bloque A)', 'Panadero (bloque B)']
SEM_OK, SEM_NO = 'Dentro de los gramajes del pliego', 'Fuera de los gramajes del pliego'
FC_OK, FC_VIG, FC_ALTO = 'En objetivo', 'Vigila el coste', 'Por encima del objetivo'

NOTA_ETG = ('AVISO FIJO. El semáforo de esta hoja solo dice si tus gramajes caben en los del pliego de la Especialidad Tradicional '
            'Garantizada (ETG): NO convierte tu pizza en «pizza napolitana». Ese nombre está reservado desde 2022 (Reglamento de '
            'Ejecución (UE) 2022/2313) y solo puede usarlo quien cumple TODO el pliego (incluido el horno de leña) y pasa la '
            'verificación oficial (Reglamento (UE) 2024/1143, arts. 72 y 73). La asociación AVPN es privada y no sustituye esa '
            'verificación. Mientras no la tengas, pon en la carta el nombre de la receta (Margherita, Marinara…).')

# ---------------------------------------------------------------------------------------------------------------------
# Datos. Precios: F1-research §2 (etiqueta «research») o estimación del kit (etiqueta «estimado»). Sin IVA.
# ---------------------------------------------------------------------------------------------------------------------
#: bloque A / B: (clave, rótulo, formato, valor por defecto, gris (mín, máx) o None, nota)
ENT_MASA = [
    ('bolas', 'Número de bolas de masa que quieres preparar', ENT, 60, None,
     'Número entero mayor que 0. Una bola, una pizza.'),
    ('peso', 'Peso de la bola (g)', GR, 250, (180, 250),
     'El pliego marca bolas de 180 a 250 g, hechas a mano. En gris, el rango del pliego.'),
    ('merma', 'Merma de amasado y boleado (%)', PCT, 0.03, None,
     'Estimación del kit: masa que se queda en la amasadora, el banco y las manos. Se suma a la masa necesaria.'),
    ('h_l', 'Harina por litro de agua (g)', GR, 1800, (1600, 1800),
     'El pliego solo fija el máximo: hasta 1.800 g de harina por litro de agua («punto de amasado»). El mínimo de 1.600 g '
     '(62,5 % de hidratación) es un cálculo del kit sobre la absorción de 55-62 % que el pliego pide a la harina.'),
    ('s_l', 'Sal por litro de agua (g)', GR, 52.5, (50, 55), 'Pliego: 50 a 55 g de sal por litro de agua.'),
    ('l_l', 'Levadura de cerveza fresca por litro de agua (g)', GR, 3, (3, 3), 'Pliego: 3 g por litro de agua.'),
    ('agua_l', 'Agua por litro (g)', GR, 1000, None, 'Constante: un litro de agua pesa 1.000 g. No se edita.'),
    ('p_h', 'Precio de la harina (€/kg, sin IVA)', EUR, 1.39, None,
     'Research: harina 00 de pizzería, 1,39 €/kg sin IVA (minorista online, octubre 2026). Pon el de tu proveedor.'),
    ('p_s', 'Precio de la sal (€/kg, sin IVA)', EUR, 0.60, None, 'Estimación del kit. Pon el de tu proveedor.'),
    ('p_l', 'Precio de la levadura fresca (€/kg, sin IVA)', EUR, 3.50, None, 'Estimación del kit (el research no encontró precio fiable). Pon el de tu proveedor.'),
]
ENT_B = [
    ('hid', 'Hidratación (agua sobre la harina, %)', PCT, 0.60, None,
     'El pliego deja entre el 55 y el 62 % según la absorción de la harina. 60 % es un ejemplo del kit.'),
    ('sal_b', 'Sal (sobre la harina, %)', '0.00%', 0.028, None, 'Pliego: entre el 2,8 y el 3,1 % sobre la harina (50-55 g por litro).'),
    ('lev_b', 'Levadura fresca (sobre la harina, %)', '0.00%', 0.002, None,
     'Pliego: aprox. 0,17 % sobre la harina (3 g por litro de agua). 0,2 % es un ejemplo del kit.'),
]

# ingredientes: nombre, €/kg sin IVA, merma %, origen del precio
ING = [
    ('Tomate pelado triturado', 2.88, 0.0, 'research: lata de 2,5 kg (Makro)'),
    ('Fior di latte', 8.82, 0.03, 'research: juliana de 2,5 kg (Makro)'),
    ('Mozzarella de búfala', 18.00, 0.03, 'estimado: el research no halló precio profesional'),
    ('Aceite de oliva virgen extra', 4.50, 0.0, 'research: garrafa de 5 l, convertido a €/kg y sin IVA'),
    ('Albahaca fresca', 12.00, 0.30, 'estimado'),
    ('Orégano seco', 15.00, 0.0, 'estimado'),
    ('Ajo', 4.00, 0.20, 'estimado'),
    ('Salami picante', 14.00, 0.03, 'estimado'),
    ('Jamón cocido', 9.50, 0.03, 'estimado'),
    ('Champiñón laminado', 3.50, 0.10, 'estimado'),
    ('Alcachofa en aceite', 9.00, 0.10, 'estimado'),
    ('Aceitunas negras', 6.50, 0.05, 'estimado'),
    ('Gorgonzola', 11.00, 0.03, 'estimado'),
    ('Grana padano', 13.00, 0.05, 'estimado'),
    ('Provola ahumada', 9.50, 0.03, 'estimado'),
    ('Salchicha fresca italiana', 8.00, 0.05, 'estimado'),
    ('Friarielli (cima de rapa)', 5.00, 0.25, 'estimado'),
    ('Tomate cherry', 3.50, 0.05, 'estimado'),
    ('Rúcula', 8.00, 0.10, 'estimado: libre para tus pizzas'),
    ('Jamón curado', 22.00, 0.05, 'estimado: libre para tus pizzas'),
    ('Nduja', 14.00, 0.02, 'estimado: libre para tus pizzas'),
    ('Ricotta', 6.00, 0.02, 'estimado: libre para tus pizzas'),
]
#: pizza → (PVP con IVA, {ingrediente: gramos})
PIZZAS = OrderedDict([
    ('Margherita', (10.00, {'Tomate pelado triturado': 70, 'Fior di latte': 90, 'Aceite de oliva virgen extra': 5, 'Albahaca fresca': 3})),
    ('Marinara', (9.50, {'Tomate pelado triturado': 100, 'Ajo': 3, 'Orégano seco': 0.5, 'Aceite de oliva virgen extra': 6})),
    ('Diavola', (12.50, {'Tomate pelado triturado': 70, 'Fior di latte': 90, 'Salami picante': 50, 'Aceite de oliva virgen extra': 5,
                         'Albahaca fresca': 2})),
    ('Capricciosa', (14.00, {'Tomate pelado triturado': 70, 'Fior di latte': 80, 'Jamón cocido': 40, 'Champiñón laminado': 40,
                             'Alcachofa en aceite': 30, 'Aceitunas negras': 15, 'Aceite de oliva virgen extra': 5})),
    ('Quattro formaggi', (13.50, {'Fior di latte': 60, 'Gorgonzola': 40, 'Grana padano': 20, 'Provola ahumada': 40,
                                  'Aceite de oliva virgen extra': 4})),
    ('Cotto e funghi', (12.50, {'Tomate pelado triturado': 70, 'Fior di latte': 90, 'Jamón cocido': 45, 'Champiñón laminado': 50,
                                'Aceite de oliva virgen extra': 5})),
    ('Bufalina', (15.50, {'Tomate pelado triturado': 60, 'Mozzarella de búfala': 125, 'Tomate cherry': 30, 'Albahaca fresca': 3,
                          'Aceite de oliva virgen extra': 5})),
    ('Salsiccia e friarielli', (14.50, {'Provola ahumada': 70, 'Salchicha fresca italiana': 80, 'Friarielli (cima de rapa)': 70,
                                        'Ajo': 2, 'Aceite de oliva virgen extra': 6})),
])
NP, NI = len(PIZZAS), len(ING)
PCOL = [openpyxl.utils.get_column_letter(5 + j) for j in range(NP)]               # E..L
ESC_IN = [('iva', 'IVA de la pizza (%)', PCT, 0.10, 'Comida preparada, en sala, para llevar y a domicilio: 10 % (Ley 37/1992, art. 91.Uno.1.1.º y 2.2.º).'),
          ('obj', 'Food cost objetivo (% sobre el precio sin IVA)', PCT, 0.25, 'Referencia habitual de pizzería: 22-30 %. Pon el que te marque tu plan financiero.'),
          ('tol', 'Margen de vigilancia sobre el objetivo (puntos)', PCT, 0.05, 'Hasta este margen por encima del objetivo, el semáforo dice «Vigila el coste»; más allá, «Por encima».'),
          ('caja', 'Coste de la caja para llevar (€, sin IVA)', EUR, 1.00, 'Research: caja de pizza de 33 cm aprox. 1 € la unidad. Pon el de tu proveedor.')]

# ---------------------------------------------------------------------------------------------------------------------
# Disposición de la hoja Masa
# ---------------------------------------------------------------------------------------------------------------------
M = OrderedDict()                                                                  # clave → fila en Masa
_f_ = 5
M_HA = _f_
for k, *_r in ENT_MASA:
    _f_ += 1
    M[k] = _f_
M_RA = _f_ + 2
SAL_A = [
    ('masa', 'Masa total que hay que preparar (g)', GR, '{bolas}*{peso}*(1+{merma})', 'Bolas × peso × (1 + merma).'),
    ('masa_l', 'Masa que sale de cada litro de agua (g)', GR, '{agua_l}+{h_l}+{s_l}+{l_l}', 'Agua + harina + sal + levadura por litro.'),
    ('litros', 'Litros de agua que necesitas', '#,##0.00', '{masa}/{masa_l}', 'Masa total ÷ masa por litro.'),
    ('agua', 'Agua (g)', GR, '{litros}*{agua_l}', ''),
    ('harina', 'Harina (g)', GR, '{litros}*{h_l}', ''),
    ('sal', 'Sal (g)', GR, '{litros}*{s_l}', ''),
    ('lev', 'Levadura fresca (g)', GR, '{litros}*{l_l}', ''),
    ('ph_h', '% panadero: hidratación (agua ÷ harina)', PCT, '{agua}/{harina}', 'Para pasarlo al bloque B.'),
    ('ph_s', '% panadero: sal (sal ÷ harina)', '0.00%', '{sal}/{harina}', ''),
    ('ph_l', '% panadero: levadura (levadura ÷ harina)', '0.00%', '{lev}/{harina}', ''),
    ('h_bola', 'Harina por bola (g)', GR, '{harina}/{bolas}', ''),
    ('bolas_kg', 'Bolas que salen de 1 kg de harina', '#,##0.0', '{bolas}*1000/{harina}', 'Con la receta del pliego, unas 6 bolas de 250 g.'),
    ('coste_a', 'Coste de la masa (€)', EUR, '({harina}*{p_h}+{sal}*{p_s}+{lev}*{p_l})/1000', 'Suma de gramos × €/kg ÷ 1.000, con la merma incluida.'),
    ('cbola_a', 'Coste de la masa por bola (€)', EUR3, '{coste_a}/{bolas}', ''),
    ('sem', 'Semáforo del pliego', 'General',
     'IF(AND({h_l}>={h_min},{h_l}<={h_max},{s_l}>={s_min},{s_l}<={s_max},{peso}>={p_min},{peso}<={p_max}),"%s","%s")' % (SEM_OK, SEM_NO),
     'Mira la harina y la sal por litro y el peso de la bola. Ojo al aviso fijo de arriba: no es un sello.'),
]
S_A = OrderedDict((k, M_RA + 1 + i) for i, (k, *_r) in enumerate(SAL_A))
M_HB = M_RA + 1 + len(SAL_A) + 1
for i, (k, *_r) in enumerate(ENT_B):
    M[k] = M_HB + 1 + i
M_RB = M_HB + 1 + len(ENT_B) + 1
SAL_B = [
    ('harina_b', 'Harina (g)', GR, '{masa}/(1+{hid}+{sal_b}+{lev_b})', 'Masa ÷ (1 + hidratación + sal + levadura), en tanto por uno.'),
    ('agua_b', 'Agua (g)', GR, '{harina_b}*{hid}', ''),
    ('sal_bg', 'Sal (g)', GR, '{harina_b}*{sal_b}', ''),
    ('lev_bg', 'Levadura fresca (g)', GR, '{harina_b}*{lev_b}', ''),
    ('coste_b', 'Coste de la masa (€)', EUR, '({harina_b}*{p_h}+{sal_bg}*{p_s}+{lev_bg}*{p_l})/1000', 'Mismos €/kg del bloque A.'),
    ('cbola_b', 'Coste de la masa por bola (€)', EUR3, '{coste_b}/{bolas}', ''),
    ('bolas_kg_b', 'Bolas que salen de 1 kg de harina', '#,##0.0', '{bolas}*1000/{harina_b}', ''),
]
S_B = OrderedDict((k, M_RB + 1 + i) for i, (k, *_r) in enumerate(SAL_B))
M_SEL = M_RB + 1 + len(SAL_B) + 1
M['metodo'] = M_SEL + 1
S_M = {'cbola': M_SEL + 2}
M_FIN = M_SEL + 3

# ---------------------------------------------------------------------------------------------------------------------
# Disposición de la hoja Escandallo
# ---------------------------------------------------------------------------------------------------------------------
E = OrderedDict((k, 5 + i) for i, (k, *_r) in enumerate(ESC_IN))
E_CM = 5 + len(ESC_IN)                                                             # coste de la masa por bola (enlace)
E_CAB = E_CM + 2
E_PVP = E_CAB + 1
E_I0 = E_PVP + 1
E_I1 = E_I0 + NI - 1
E_R = E_I1 + 2
RES = [
    ('peso_t', 'Peso en crudo de la pizza (g)', GR),
    ('c_masa', 'Coste de la masa (€)', EUR3),
    ('c_ing', 'Coste de los ingredientes (€)', EUR3),
    ('c_tot', 'COSTE DE LA PIZZA (€)', EUR3),
    ('pvp_sin', 'Precio de venta sin IVA (€)', EUR),
    ('fc', 'FOOD COST (%)', PCT),
    ('sem', 'Semáforo contra el objetivo', 'General'),
    ('margen', 'Margen por pizza (€, precio sin IVA − coste)', EUR),
    ('pvp_obj', 'PVP con IVA para llegar al objetivo (€)', EUR),
    ('c_llevar', 'Coste para llevar: pizza + caja (€)', EUR3),
    ('fc_llevar', 'Food cost para llevar con la caja (%)', PCT),
]
R = OrderedDict((k, E_R + 1 + i) for i, (k, *_r) in enumerate(RES))
E_FIN = E_R + len(RES) + 1


def _ming(nombre):
    return [i for i, (n, *_r) in enumerate(ING) if n == nombre][0]


def valores_defecto():
    """Todas las entradas por defecto en un solo dict plano."""
    v = OrderedDict()
    for k, _r, _f2, val, _g, _n in ENT_MASA + ENT_B:
        v[k] = val
    v['metodo'] = METODOS[0]
    for k, _r, _f2, val, _n in ESC_IN:
        v[k] = val
    for i, (nom, p, m, _o) in enumerate(ING):
        v['p%d' % i], v['m%d' % i] = p, m
    for j, (nom, (pvp, gr)) in enumerate(PIZZAS.items()):
        v['pvp%d' % j] = pvp
        for i, (ni, *_r) in enumerate(ING):
            v['g%d_%d' % (i, j)] = float(gr.get(ni, 0))
    return v


# ---------------------------------------------------------------------------------------------------------------------
# Cálculo en Python (espejo de las fórmulas) — fuente del autotest
# ---------------------------------------------------------------------------------------------------------------------
GRIS_MM = dict(h_min=1600, h_max=1800, s_min=50, s_max=55, p_min=180, p_max=250)


def calcular(v):
    r = OrderedDict()
    r['masa'] = v['bolas'] * v['peso'] * (1 + v['merma'])
    r['masa_l'] = v['agua_l'] + v['h_l'] + v['s_l'] + v['l_l']
    r['litros'] = r['masa'] / r['masa_l']
    r['agua'] = r['litros'] * v['agua_l']
    r['harina'] = r['litros'] * v['h_l']
    r['sal'] = r['litros'] * v['s_l']
    r['lev'] = r['litros'] * v['l_l']
    r['ph_h'] = r['agua'] / r['harina']
    r['ph_s'] = r['sal'] / r['harina']
    r['ph_l'] = r['lev'] / r['harina']
    r['h_bola'] = r['harina'] / v['bolas']
    r['bolas_kg'] = v['bolas'] * 1000 / r['harina']
    r['coste_a'] = (r['harina'] * v['p_h'] + r['sal'] * v['p_s'] + r['lev'] * v['p_l']) / 1000
    r['cbola_a'] = r['coste_a'] / v['bolas']
    g = GRIS_MM
    r['sem'] = SEM_OK if (g['h_min'] <= v['h_l'] <= g['h_max'] and g['s_min'] <= v['s_l'] <= g['s_max']
                          and g['p_min'] <= v['peso'] <= g['p_max']) else SEM_NO
    r['harina_b'] = r['masa'] / (1 + v['hid'] + v['sal_b'] + v['lev_b'])
    r['agua_b'] = r['harina_b'] * v['hid']
    r['sal_bg'] = r['harina_b'] * v['sal_b']
    r['lev_bg'] = r['harina_b'] * v['lev_b']
    r['coste_b'] = (r['harina_b'] * v['p_h'] + r['sal_bg'] * v['p_s'] + r['lev_bg'] * v['p_l']) / 1000
    r['cbola_b'] = r['coste_b'] / v['bolas']
    r['bolas_kg_b'] = v['bolas'] * 1000 / r['harina_b']
    r['cbola'] = r['cbola_a'] if v['metodo'] == METODOS[0] else r['cbola_b']
    neto = [v['p%d' % i] / (1 - v['m%d' % i]) for i in range(NI)]
    r['neto'] = neto
    for j in range(NP):
        gs = [v['g%d_%d' % (i, j)] for i in range(NI)]
        p = OrderedDict()
        p['peso_t'] = v['peso'] + sum(gs)
        p['c_masa'] = r['cbola']
        p['c_ing'] = sum(g_ * n for g_, n in zip(gs, neto)) / 1000
        p['c_tot'] = p['c_masa'] + p['c_ing']
        p['pvp_sin'] = v['pvp%d' % j] / (1 + v['iva'])
        p['fc'] = p['c_tot'] / p['pvp_sin']
        p['sem'] = FC_OK if p['fc'] <= v['obj'] else (FC_VIG if p['fc'] <= v['obj'] + v['tol'] else FC_ALTO)
        p['margen'] = p['pvp_sin'] - p['c_tot']
        p['pvp_obj'] = p['c_tot'] / v['obj'] * (1 + v['iva'])
        p['c_llevar'] = p['c_tot'] + v['caja']
        p['fc_llevar'] = p['c_llevar'] / p['pvp_sin']
        r['p%d' % j] = p
    return r


# ---------------------------------------------------------------------------------------------------------------------
# Helpers (de la finca)
# ---------------------------------------------------------------------------------------------------------------------
def _f(ws, coord, v, b=False, sz=11, i=False, fmt=None, wrap=False, fill=None, locked=True):
    c = ws[coord]
    c.value = v
    c.font = Font(name=FUENTE, size=sz, bold=b, italic=i)
    if fmt:
        c.number_format = fmt
    c.alignment = Alignment(wrap_text=wrap, vertical='top')
    if fill:
        c.fill = fill
    c.protection = Protection(locked=locked)
    return c


def _nota(ws, fila, texto, desde='E', hasta='F', chars=80.0):
    if not texto:
        return
    _f(ws, '%s%d' % (desde, fila), texto, wrap=True)
    ws.merge_cells('%s%d:%s%d' % (desde, fila, hasta, fila))
    ws.row_dimensions[fila].height = max(16, 15 * math.ceil(len(texto) / chars))


def _imprimir(ws, horizontal=False):
    ws.page_setup.paperSize = 9                                                   # A4
    ws.page_setup.orientation = 'landscape' if horizontal else 'portrait'
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.oddFooter.center.text = 'aichef.pro/%s · página &P de &N' % PID
    ws.print_options.horizontalCentered = True
    ws.protection.sheet = True                                                    # sin contraseña
    ws.sheet_view.showGridLines = False


def _dv(ws, tipo, celdas, f1=None, f2=None, op=None, titulo='Valor no válido', msg=''):
    dv = DataValidation(type=tipo, formula1=f1, formula2=f2, operator=op, allow_blank=False, showErrorMessage=True,
                        errorTitle=titulo, error=msg)
    ws.add_data_validation(dv)
    for c in celdas:
        dv.add(c)


def a_excel(expr, ref):
    return re.sub(r'\{(\w+)\}', lambda m: ref[m.group(1)], expr)


def refs_masa():
    ref = dict((k, '$B$%d' % f) for k, f in M.items())
    ref.update((k, '$B$%d' % f) for k, f in S_A.items())
    ref.update((k, '$B$%d' % f) for k, f in S_B.items())
    for k, _r, _f2, _v, g, _n in ENT_MASA:
        if g:
            ref[k[0] + '_min'] = '$C$%d' % M[k]
            ref[k[0] + '_max'] = '$D$%d' % M[k]
    ref['p_min'], ref['p_max'] = '$C$%d' % M['peso'], '$D$%d' % M['peso']
    ref['h_min'], ref['h_max'] = '$C$%d' % M['h_l'], '$D$%d' % M['h_l']
    ref['s_min'], ref['s_max'] = '$C$%d' % M['s_l'], '$D$%d' % M['s_l']
    return ref


# ---------------------------------------------------------------------------------------------------------------------
# Hojas
# ---------------------------------------------------------------------------------------------------------------------
def hoja_masa(wb):
    ws = wb.active
    ws.title = HM
    for col, w in (('A', 56), ('B', 17), ('C', 12), ('D', 12), ('E', 48), ('F', 30)):
        ws.column_dimensions[col].width = w
    _f(ws, 'A1', 'CALCULADORA DE MASA PARA PIZZA — A LOS GRAMAJES DEL PLIEGO Y POR % PANADERO', b=True, sz=14)
    _f(ws, 'A2', 'Escribe solo en las celdas VERDES: todo lo demás se calcula. Importes sin IVA. Los precios de los ingredientes '
                 'son ejemplos (research de octubre 2026 o estimación del kit): cámbialos por los de tu proveedor.', wrap=True)
    ws.merge_cells('A2:F2')
    ws.row_dimensions[2].height = 32
    _f(ws, 'A3', NOTA_ETG, wrap=True, b=True, fill=GRIS)
    ws.merge_cells('A3:F3')
    ws.row_dimensions[3].height = 92
    for col, t in zip('ABCDE', ('A. POR LITRO DE AGUA (gramajes del pliego)', 'Valor', 'Pliego mín.', 'Pliego máx.', 'Notas')):
        _f(ws, '%s%d' % (col, M_HA), t, b=True, fill=GRIS, wrap=True)
    for k, rot, fmt, val, gris, nota in ENT_MASA:
        f = M[k]
        _f(ws, 'A%d' % f, rot)
        if k == 'agua_l':
            _f(ws, 'B%d' % f, val, fmt=fmt, fill=GRIS)
        else:
            _f(ws, 'B%d' % f, val, fmt=fmt, fill=VERDE, locked=False)
        if gris:
            _f(ws, 'C%d' % f, gris[0], fmt=fmt, fill=GRIS)
            _f(ws, 'D%d' % f, gris[1], fmt=fmt, fill=GRIS)
        _nota(ws, f, nota)
    ref = refs_masa()
    for col, t in zip('ABE', ('RESULTADO DEL BLOQUE A', 'Valor', 'Notas')):
        _f(ws, '%s%d' % (col, M_RA), t, b=True, fill=GRIS)
    for k, rot, fmt, expr, nota in SAL_A:
        f = S_A[k]
        neg = k in ('masa', 'cbola_a', 'sem')
        _f(ws, 'A%d' % f, rot, b=neg)
        _f(ws, 'B%d' % f, '=' + a_excel(expr, ref), fmt=fmt, b=neg, wrap=(k == 'sem'))
        _nota(ws, f, nota)
    ws.row_dimensions[S_A['sem']].height = 48
    for col, t in zip('ABE', ('B. POR % PANADERO (sobre la harina)', 'Valor', 'Notas')):
        _f(ws, '%s%d' % (col, M_HB), t, b=True, fill=GRIS)
    for k, rot, fmt, val, gris, nota in ENT_B:
        f = M[k]
        _f(ws, 'A%d' % f, rot)
        _f(ws, 'B%d' % f, val, fmt=fmt, fill=VERDE, locked=False)
        _nota(ws, f, nota)
    for col, t in zip('ABE', ('RESULTADO DEL BLOQUE B (misma masa total que el A)', 'Valor', 'Notas')):
        _f(ws, '%s%d' % (col, M_RB), t, b=True, fill=GRIS)
    for k, rot, fmt, expr, nota in SAL_B:
        f = S_B[k]
        neg = k == 'cbola_b'
        _f(ws, 'A%d' % f, rot, b=neg)
        _f(ws, 'B%d' % f, '=' + a_excel(expr, ref), fmt=fmt, b=neg)
        _nota(ws, f, nota)
    _f(ws, 'A%d' % M_SEL, 'COSTE DE LA MASA QUE USA EL ESCANDALLO', b=True, fill=GRIS)
    _f(ws, 'B%d' % M_SEL, 'Valor', b=True, fill=GRIS)
    _f(ws, 'E%d' % M_SEL, 'Notas', b=True, fill=GRIS)
    f = M['metodo']
    _f(ws, 'A%d' % f, 'Bloque que alimenta el escandallo')
    _f(ws, 'B%d' % f, METODOS[0], fill=VERDE, locked=False, wrap=True)
    ws.row_dimensions[f].height = 32
    _nota(ws, f, 'Elige qué coste de masa por bola usa la hoja «Escandallo»: el del pliego (A) o el de tus porcentajes (B).')
    f = S_M['cbola']
    _f(ws, 'A%d' % f, 'Coste de la masa por bola para el escandallo (€)', b=True)
    _f(ws, 'B%d' % f, '=IF(%s="%s",%s,%s)' % ('$B$%d' % M['metodo'], METODOS[0], '$B$%d' % S_A['cbola_a'], '$B$%d' % S_B['cbola_b']),
       fmt=EUR3, b=True)
    # semáforo en color
    sem = '$B$%d' % S_A['sem']
    ws.conditional_formatting.add('B%d' % S_A['sem'], FormulaRule(formula=['%s="%s"' % (sem, SEM_OK)], fill=OK))
    ws.conditional_formatting.add('B%d' % S_A['sem'], FormulaRule(formula=['%s="%s"' % (sem, SEM_NO)], fill=ROJO))
    # validaciones
    cel = lambda k: 'B%d' % M[k]
    _dv(ws, 'whole', [cel('bolas')], f1='1', f2='100000', op='between', titulo='Número no válido',
        msg='Escribe un número entero de 1 o más.')
    _dv(ws, 'decimal', [cel(k) for k in ('peso', 'h_l', 's_l', 'l_l', 'p_h', 'p_s', 'p_l')], f1='0', op='greaterThan',
        titulo='Valor no válido', msg='Escribe un número mayor que 0.')
    _dv(ws, 'decimal', [cel('merma'), cel('sal_b'), cel('lev_b')], f1='0', f2='1', op='between', titulo='Porcentaje no válido',
        msg='Escribe un porcentaje entre 0 y 1 (0,03 = 3 %).')
    _dv(ws, 'decimal', [cel('hid')], f1='0', f2='2', op='between', titulo='Hidratación no válida',
        msg='Escribe un porcentaje entre 0 y 200 % (0,6 = 60 %).')
    _dv(ws, 'list', [cel('metodo')], f1='"%s"' % ','.join(METODOS), titulo='Opción no válida', msg='Elige una opción de la lista.')
    _imprimir(ws)


def hoja_escandallo(wb):
    ws = wb.create_sheet(HE)
    for col, w in (('A', 40), ('B', 13), ('C', 10), ('D', 13)):
        ws.column_dimensions[col].width = w
    for col in PCOL:
        ws.column_dimensions[col].width = 14
    ws.column_dimensions['M'].width = 46
    _f(ws, 'A1', 'ESCANDALLO POR PIZZA — 8 PIZZAS DE EJEMPLO, EN GRAMOS', b=True, sz=14)
    _f(ws, 'A2', 'Escribe solo en las celdas VERDES (ingredientes, precios, mermas, gramos, nombres y precios de carta). Importes sin IVA '
                 'salvo el precio de carta. Los precios son ejemplos: pon los de tu proveedor. La carta de ejemplo usa el nombre de la '
                 'receta, no «napolitana» (ver el aviso de la hoja «%s»).' % HM, wrap=True)
    ws.merge_cells('A2:M2')
    ws.row_dimensions[2].height = 32
    for col, t in zip('ABE', ('DATOS GENERALES', 'Valor', 'Notas')):
        _f(ws, '%s4' % col, t, b=True, fill=GRIS)
    for k, rot, fmt, val, nota in ESC_IN:
        f = E[k]
        _f(ws, 'A%d' % f, rot)
        _f(ws, 'B%d' % f, val, fmt=fmt, fill=VERDE, locked=False)
        _nota(ws, f, nota, desde='E', hasta='M', chars=140.0)
    _f(ws, 'A%d' % E_CM, 'Coste de la masa por bola (de la hoja «%s»)' % HM, b=True)
    _f(ws, 'B%d' % E_CM, "='%s'!$B$%d" % (HM, S_M['cbola']), fmt=EUR3, b=True)
    _nota(ws, E_CM, 'Se calcula en la hoja «%s» (bloque A o B, según elijas allí).' % HM, desde='E', hasta='M', chars=140.0)
    # cabecera de la tabla
    for col, t in (('A', 'INGREDIENTE'), ('B', '€/kg sin IVA'), ('C', 'Merma'), ('D', '€/kg neto')):
        _f(ws, '%s%d' % (col, E_CAB), t, b=True, fill=GRIS, wrap=True)
    _f(ws, 'M%d' % E_CAB, 'ORIGEN DEL PRECIO DE EJEMPLO', b=True, fill=GRIS, wrap=True)
    _f(ws, 'A%d' % E_PVP, 'Precio de carta con IVA (€)', b=True)
    for j, (nom, (pvp, _gr)) in enumerate(PIZZAS.items()):
        col = PCOL[j]
        _f(ws, '%s%d' % (col, E_CAB), nom, b=True, fill=VERDE, wrap=True, locked=False)
        _f(ws, '%s%d' % (col, E_PVP), pvp, fmt=EUR, fill=VERDE, locked=False, b=True)
    ws.row_dimensions[E_CAB].height = 32
    for i, (nom, p, m, origen) in enumerate(ING):
        f = E_I0 + i
        _f(ws, 'A%d' % f, nom, fill=VERDE, locked=False)
        _f(ws, 'B%d' % f, p, fmt=EUR, fill=VERDE, locked=False)
        _f(ws, 'C%d' % f, m, fmt='0%', fill=VERDE, locked=False)
        _f(ws, 'D%d' % f, '=B%d/(1-C%d)' % (f, f), fmt=EUR)
        for j, (pn, (_pvp, gr)) in enumerate(PIZZAS.items()):
            _f(ws, '%s%d' % (PCOL[j], f), float(gr.get(nom, 0)), fmt='#,##0.0;-#,##0.0;"-"', fill=VERDE, locked=False)
        _f(ws, 'M%d' % f, origen, i=True)
    _f(ws, 'A%d' % (E_R), 'RESULTADO POR PIZZA (gramos en las columnas de arriba)', b=True, fill=GRIS)
    for j in range(NP):
        _f(ws, '%s%d' % (PCOL[j], E_R), '=%s%d' % (PCOL[j], E_CAB), b=True, fill=GRIS, wrap=True)
    ws.row_dimensions[E_R].height = 32
    for k, rot, fmt in RES:
        f = R[k]
        neg = k in ('c_tot', 'fc')
        _f(ws, 'A%d' % f, rot, b=neg)
        for j in range(NP):
            c = PCOL[j]
            rng = '%s$%d:%s$%d' % (c, E_I0, c, E_I1)
            ex = {
                'peso_t': "'%s'!$B$%d+SUM(%s)" % (HM, M['peso'], rng),
                'c_masa': '$B$%d' % E_CM,
                'c_ing': 'SUMPRODUCT(%s,$D$%d:$D$%d)/1000' % (rng, E_I0, E_I1),
                'c_tot': '%s%d+%s%d' % (c, R['c_masa'], c, R['c_ing']),
                'pvp_sin': '%s%d/(1+$B$%d)' % (c, E_PVP, E['iva']),
                'fc': '%s%d/%s%d' % (c, R['c_tot'], c, R['pvp_sin']),
                'sem': 'IF(%s%d<=$B$%d,"%s",IF(%s%d<=$B$%d+$B$%d,"%s","%s"))' % (
                    c, R['fc'], E['obj'], FC_OK, c, R['fc'], E['obj'], E['tol'], FC_VIG, FC_ALTO),
                'margen': '%s%d-%s%d' % (c, R['pvp_sin'], c, R['c_tot']),
                'pvp_obj': '%s%d/$B$%d*(1+$B$%d)' % (c, R['c_tot'], E['obj'], E['iva']),
                'c_llevar': '%s%d+$B$%d' % (c, R['c_tot'], E['caja']),
                'fc_llevar': '%s%d/%s%d' % (c, R['c_llevar'], c, R['pvp_sin']),
            }[k]
            _f(ws, '%s%d' % (c, f), '=' + ex, fmt=fmt, b=neg, wrap=(k == 'sem'))
    ws.row_dimensions[R['sem']].height = 32
    rng = '%s%d:%s%d' % (PCOL[0], R['sem'], PCOL[-1], R['sem'])
    top = '%s%d' % (PCOL[0], R['sem'])
    for txt, fill in ((FC_OK, OK), (FC_VIG, AMBAR), (FC_ALTO, ROJO)):
        ws.conditional_formatting.add(rng, FormulaRule(formula=['%s="%s"' % (top, txt)], fill=fill))
    f = E_FIN
    _f(ws, 'A%d' % f, 'El food cost de la tabla es el de la pizza servida en sala. Para llevar se suma la caja al coste y el precio '
                      'de carta es el mismo; si cobras más por llevar o a domicilio, cambia el precio de la pizza para ver el efecto. '
                      'No incluye la comisión de la plataforma de reparto, que se descuenta del precio, no del coste de la pizza.',
       i=True, wrap=True)
    ws.merge_cells('A%d:M%d' % (f, f))
    ws.row_dimensions[f].height = 34
    # validaciones
    pc = lambda cols, f1_, f2_: ['%s%d' % (c, f1_) for c in cols]
    _dv(ws, 'decimal', ['B%d' % E['iva'], 'B%d' % E['obj'], 'B%d' % E['tol']] + ['C%d' % (E_I0 + i) for i in range(NI)],
        f1='0', f2='0.9', op='between', titulo='Porcentaje no válido', msg='Escribe un porcentaje entre 0 y 0,9 (0,1 = 10 %).')
    _dv(ws, 'decimal', ['B%d' % E['caja']] + ['B%d' % (E_I0 + i) for i in range(NI)]
        + ['%s%d' % (c, E_PVP) for c in PCOL] + ['%s%d' % (c, E_I0 + i) for c in PCOL for i in range(NI)],
        f1='0', op='greaterThanOrEqual', titulo='Importe no válido', msg='Escribe un número mayor o igual que 0.')
    _dv(ws, 'decimal', ['B%d' % E['obj']], f1='0.01', f2='0.9', op='between', titulo='Objetivo no válido',
        msg='El objetivo no puede ser 0 (se divide entre él): escribe entre 0,01 y 0,9.')
    ws.page_setup.paperSize = 9
    _imprimir(ws, horizontal=True)


INSTRUCCIONES = [
    ('INSTRUCCIONES DE USO — Calculadora de masa y escandallo por pizza', True),
    ('Dos herramientas en un libro: la hoja «Masa» te dice cuánta harina, agua, sal y levadura necesitas para el número de bolas '
     'que vas a preparar y cuánto cuesta cada bola; la hoja «Escandallo» calcula el coste de 8 pizzas de ejemplo, su food cost, '
     'su margen y el precio de carta que necesitas. Solo se escribe en las celdas VERDES; lo demás se calcula.', False),
    ('AVISO SOBRE EL NOMBRE', True),
    (NOTA_ETG, False),
    ('HOJA «MASA»', True),
    ('1. Bloque A (por litro de agua): escribe cuántas bolas quieres, su peso y la merma de amasado. Las cantidades por litro vienen '
     'del pliego de la Especialidad Tradicional Garantizada (Reglamento (UE) n.º 97/2010, anexo II): hasta 1.800 g de harina, 50 a 55 g '
     'de sal y 3 g de levadura de cerveza por litro de agua, con bolas de 180 a 250 g. El rango del pliego aparece en gris al lado.', False),
    ('2. El semáforo del bloque A dice si la harina, la sal y el peso de la bola caben en esos rangos. Verde no significa «pizza '
     'napolitana»: el pliego exige además fermentación en ciclo continuo, formado a mano y horno de leña, y la verificación oficial.', False),
    ('3. Bloque B (por % panadero): escribe la hidratación, la sal y la levadura como porcentaje del peso de la harina, que es como '
     'trabaja el oficio. Usa la misma masa total que el bloque A. Si le pones los porcentajes que el bloque A calcula, obtienes lo mismo.', False),
    ('4. Abajo eliges qué coste de masa por bola alimenta el escandallo: el del pliego (A) o el de tus porcentajes (B).', False),
    ('HOJA «ESCANDALLO»', True),
    ('1. Escribe el IVA (10 %), el food cost objetivo, el margen de vigilancia y el coste de la caja para llevar.', False),
    ('2. En la tabla, cada fila es un ingrediente con su precio por kilo sin IVA y su merma (la parte que se pierde al limpiar o '
     'recortar). El precio neto por kilo divide el precio entre (1 − merma). Los gramos de cada pizza van en su columna; puedes cambiar '
     'los nombres, los gramos y los precios de carta, y usar las filas libres para tus propios ingredientes.', False),
    ('3. Para cada pizza ves el coste (masa + ingredientes), el precio sin IVA, el food cost con semáforo, el margen por pizza, el precio '
     'de carta con IVA que necesitarías para llegar a tu objetivo y el food cost para llevar con la caja.', False),
    ('4. La Margherita de ejemplo queda aproximadamente un 15 % de food cost. Con búfala, embutidos italianos o fiambres, sube. Si una pizza '
     'sale por encima del objetivo, sube el precio, baja el gramaje o cambia un ingrediente.', False),
    ('PRECIOS Y FUENTES', True),
    ('Los precios por defecto son ejemplos. Los marcados como «research» salen del estudio de mercado del kit (octubre 2026: harina '
     'de pizzería, fior di latte y tomate pelado de catálogos de mayoristas, aceite de oferta, caja de 33 cm); el resto son estimaciones '
     'del kit. Cámbialos por los de tu proveedor antes de decidir precios. El precio de carta de ejemplo (9,50 a 15,50 €, Margherita a 10 €) '
     'es el del caso del plan de negocio.', False),
    ('El IVA de la comida preparada para consumir en el local, para llevar o a domicilio es el 10 % (Ley 37/1992, art. 91.Uno.1.1.º y '
     '2.2.º). Las bebidas alcohólicas y los refrescos con azúcares o edulcorantes añadidos que salen del local llevan el 21 % y no '
     'están en este escandallo.', False),
    ('Fuentes: pliego de condiciones de la ETG «Pizza Napoletana» (Reglamento (UE) n.º 97/2010, anexo II); Reglamento de Ejecución (UE) '
     '2022/2313; Reglamento (UE) 2024/1143, arts. 68 a 73; Ley 37/1992 del IVA, art. 91; research del producto (octubre 2026).', False),
]


def hoja_instrucciones(wb):
    ws = wb.create_sheet(HI)
    ws.column_dimensions['A'].width = 110
    f = 1
    for t, b in INSTRUCCIONES:
        _f(ws, 'A%d' % f, t, b=b, sz=13 if f == 1 else 11, wrap=True)
        if not b:
            ws.row_dimensions[f].height = max(16, 15 * math.ceil(len(t) / 120.0))
        f += 2 if b and f > 1 else 1
    f += 1
    for t in ('Para editar una celda que no esté en verde: Revisar → Desproteger hoja (no tiene contraseña).',
              'Más plantillas y kits del catálogo en aichef.pro/productos-digitales', motor.BIO_LINE,
              motor.VERSION_FMT.format(pid=PID)):
        _f(ws, 'A%d' % f, t, wrap=True)
        f += 1
    _imprimir(ws)


def generar(destino):
    wb = openpyxl.Workbook()
    hoja_masa(wb)
    hoja_escandallo(wb)
    hoja_instrucciones(wb)
    pr = wb.properties
    pr.title = '%s · v%s' % (TITULO, motor.VERSION)
    pr.subject = 'Plan de Negocio: Pizzería Napolitana'
    pr.description = 'aichef.pro/%s' % PID
    pr.keywords = 'plan negocio pizzería, calculadora masa pizza, escandallo pizza, AI Chef Pro'
    pr.creator = pr.lastModifiedBy = motor.CREATOR
    os.makedirs(os.path.dirname(destino), exist_ok=True)
    wb.save(destino)
    sys.path.insert(0, PD)
    import inject_cache as IC                                                      # noqa: E402
    nf, inj, fail = IC.inject(destino)
    print('%s: formulas=%d cache_inyectado=%d fallos_pycel=%d' % (os.path.basename(destino), nf, inj, fail))
    return fail


# ---------------------------------------------------------------------------------------------------------------------
# Autotest
# ---------------------------------------------------------------------------------------------------------------------
PROHIBIDOS = [r'\bfreelance', r'validad[oa]s?', r'Espa[ñn]a,? 2026', r'carnet', r'carné', r'tarifa plana', r'RD 393', r'certificad',
              r'\bverace\b', r'service charge|sales tax']
CELDA = {}                                                                         # clave de entrada → (hoja, celda)
for _k, _f2 in M.items():
    CELDA[_k] = (HM, 'B%d' % _f2)
for _k, _f2 in E.items():
    CELDA[_k] = (HE, 'B%d' % _f2)
for _i in range(NI):
    CELDA['p%d' % _i] = (HE, 'B%d' % (E_I0 + _i))
    CELDA['m%d' % _i] = (HE, 'C%d' % (E_I0 + _i))
    for _j in range(NP):
        CELDA['g%d_%d' % (_i, _j)] = (HE, '%s%d' % (PCOL[_j], E_I0 + _i))
for _j in range(NP):
    CELDA['pvp%d' % _j] = (HE, '%s%d' % (PCOL[_j], E_PVP))


def celdas_salida(r):
    """(hoja, celda, valor esperado, etiqueta) de todas las salidas."""
    out = []
    for k in S_A:
        out.append((HM, 'B%d' % S_A[k], r[k], k))
    for k in S_B:
        out.append((HM, 'B%d' % S_B[k], r[k], k))
    out.append((HM, 'B%d' % S_M['cbola'], r['cbola'], 'cbola'))
    for i in range(NI):
        out.append((HE, 'D%d' % (E_I0 + i), r['neto'][i], 'neto%d' % i))
    for j in range(NP):
        for k in R:
            out.append((HE, '%s%d' % (PCOL[j], R[k]), r['p%d' % j][k], '%s[%d]' % (k, j)))
    return out


def _distinto(x, v):
    if isinstance(v, str):
        return x != v
    return not isinstance(x, (int, float)) or abs(float(x) - v) > 0.01


def autotest(destino):
    from pycel import ExcelCompiler
    err, n = [], 0
    wv = openpyxl.load_workbook(destino, data_only=True)
    wf = openpyxl.load_workbook(destino)
    base = valores_defecto()
    j2 = dict(base, bolas=120, peso=200, merma=0.05, h_l=1700, s_l=55, l_l=3, p_h=1.60, hid=0.62, sal_b=0.03, lev_b=0.0017,
              metodo=METODOS[1], iva=0.10, obj=0.28, tol=0.03, caja=1.2, pvp0=11.0, pvp6=16.0, p1=9.5, m1=0.05, g0_2=60.0, g7_5=100.0)
    j3 = dict(base, bolas=24, peso=300, merma=0.0, h_l=1500, s_l=60, p_s=0.9, p_l=4.0, hid=0.55, obj=0.2, tol=0.1, caja=0.8,
              pvp3=17.0, p2=20.0, m2=0.1, g1_1=75.0, g3_7=40.0)
    # 1. caché y pycel del juego por defecto; pycel de los otros dos con set_value
    xl = ExcelCompiler(filename=destino)
    r0 = calcular(base)
    for hoja, c, v, lab in celdas_salida(r0):
        n += 1
        for et, x in (('pycel', xl.evaluate("'%s'!%s" % (hoja, c))), ('caché', wv[hoja][c].value)):
            if _distinto(x, v):
                err.append('defecto · %s!%s (%s): %s %r ≠ esperado %r' % (hoja, c, lab, et, x, v))
    for nombre, jj in (('juego 2', j2), ('juego 3', j3)):
        r = calcular(jj)
        xl2 = ExcelCompiler(filename=destino)
        for hoja, c, v, lab in celdas_salida(r0):
            xl2.evaluate("'%s'!%s" % (hoja, c))
        for k, v in jj.items():
            if v != base[k]:
                h, c = CELDA[k]
                xl2.set_value("'%s'!%s" % (h, c), v)
        for hoja, c, v, lab in celdas_salida(r):
            n += 1
            x = xl2.evaluate("'%s'!%s" % (hoja, c))
            if _distinto(x, v):
                err.append('%s · %s!%s (%s): pycel %r ≠ esperado %r' % (nombre, hoja, c, lab, x, v))
    # 2. Σ ingredientes = masa (A y B) y B con los % de A reproduce A
    for nombre, jj in (('defecto', base), ('juego 2', j2), ('juego 3', j3)):
        r = calcular(jj)
        if abs(r['agua'] + r['harina'] + r['sal'] + r['lev'] - r['masa']) > 0.5:
            err.append('%s: Σ ingredientes del A ≠ masa' % nombre)
        if abs(r['harina_b'] + r['agua_b'] + r['sal_bg'] + r['lev_bg'] - r['masa']) > 0.5:
            err.append('%s: Σ ingredientes del B ≠ masa' % nombre)
        rb = calcular(dict(jj, hid=r['ph_h'], sal_b=r['ph_s'], lev_b=r['ph_l']))
        for a, b in (('harina', 'harina_b'), ('agua', 'agua_b'), ('sal', 'sal_bg'), ('lev', 'lev_bg'), ('coste_a', 'coste_b'),
                     ('cbola_a', 'cbola_b')):
            n += 1
            if abs(rb[a] - rb[b]) > 0.01:
                err.append('%s: el B con los %% del A no reproduce el A (%s %.4f ≠ %s %.4f)' % (nombre, a, rb[a], b, rb[b]))
    xl3 = ExcelCompiler(filename=destino)                                       # lo mismo, con pycel sobre la hoja
    for k in S_B:
        xl3.evaluate("'%s'!B%d" % (HM, S_B[k]))
    for k, ck in (('hid', 'ph_h'), ('sal_b', 'ph_s'), ('lev_b', 'ph_l')):
        xl3.set_value("'%s'!B%d" % (HM, M[k]), r0[ck])
    for a, b in (('harina', 'harina_b'), ('agua', 'agua_b'), ('sal', 'sal_bg'), ('lev', 'lev_bg'), ('cbola_a', 'cbola_b')):
        n += 1
        xa, xb = xl3.evaluate("'%s'!B%d" % (HM, S_A[a])), xl3.evaluate("'%s'!B%d" % (HM, S_B[b]))
        if abs(xa - xb) > 0.01:
            err.append('pycel: el B con los %% del A no reproduce el A (%s %.4f ≠ %s %.4f)' % (a, xa, b, xb))
    # 3. valores con sentido
    marg = r0['p0']
    if not 0.12 <= marg['fc'] <= 0.30:
        err.append('food cost de la Margherita %.1f %% fuera de 12-30 %%' % (marg['fc'] * 100))
    if r0['sem'] != SEM_OK:
        err.append('el valor por defecto no cae dentro del pliego')
    if calcular(dict(base, h_l=2000))['sem'] != SEM_NO or calcular(dict(base, peso=300))['sem'] != SEM_NO:
        err.append('el semáforo no se pone rojo fuera del pliego')
    pizzas = list(PIZZAS)
    if any('napolit' in p.lower() for p in pizzas):
        err.append('una pizza de la carta se llama «napolitana»')
    pv = [v[0] for v in PIZZAS.values()]
    if not 9.5 <= min(pv) or max(pv) > 15.5 or PIZZAS['Margherita'][0] != 10.0:
        err.append('carta fuera de 9,50-15,50 € o Margherita ≠ 10 €')
    # 4. formato, protección, DV, fórmulas
    if motor.gate_formatos(wf, FICHERO):
        err.append('gate_formatos: %r' % motor.gate_formatos(wf, FICHERO)[:3])
    ws = wf[HM]
    dvs = dict((str(rng), dv) for dv in wf[HM].data_validations.dataValidation for rng in dv.sqref.ranges)
    if len(wf[HE].data_validations.dataValidation) < 3:
        err.append('faltan DV en Escandallo')
    dv = dvs.get('B%d' % M['bolas'])
    if dv is None or dv.type != 'whole' or str(dv.formula1) != '1':
        err.append('DV del número de bolas (entero > 0) ausente o incorrecta')
    dv = dvs.get('B%d' % M['metodo'])
    if dv is None or dv.type != 'list':
        err.append('DV de lista del método ausente')
    for k in ('peso', 'h_l', 's_l', 'p_h'):
        if dvs.get('B%d' % M[k]) is None:
            err.append('DV ausente en %s' % k)
    textos = []
    for w in wf.worksheets:
        if w.page_setup.paperSize != 9 or not w.protection.sheet:
            err.append('%s: A4 y protección' % w.title)
        for row in w.iter_rows():
            for c in row:
                v = c.value
                if isinstance(v, str) and c.data_type != 'f':
                    textos.append(v)
                if c.data_type == 'f':
                    if wv[w.title][c.coordinate].value is None:
                        err.append('%s!%s: fórmula sin caché' % (w.title, c.coordinate))
                    sin_refs = re.sub(r'"[^"]*"', '', re.sub(r"'[^']+'!|\$?[A-Z]{1,2}\$?\d+", '', str(c.value)))
                    for num in re.findall(r'(?<![A-Za-z])\d+\.?\d*', sin_refs):
                        if num not in ('0', '1', '1000'):
                            err.append('%s!%s: número tecleado %s en %s' % (w.title, c.coordinate, num, c.value))
                if c.fill.fgColor.rgb in ('00E8F5E9', 'FFE8F5E9') and c.protection.locked:
                    err.append('%s!%s: verde bloqueada' % (w.title, c.coordinate))
    todo = '\n'.join(textos)
    for rx in PROHIBIDOS:
        if re.search(rx, todo, re.I):
            err.append('prohibido /%s/' % rx)
    if re.search(r'[^\x00-ɏ‐-‧€−→·]', todo):
        err.append('caracteres no latinos o raros: %r' % sorted(set(re.findall(r'[^\x00-ɏ‐-‧€−→·]', todo))))
    if 'NO convierte tu pizza en «pizza napolitana»' not in todo:
        err.append('falta el aviso fijo de la ETG')
    if motor.VERSION_FMT.format(pid=PID) not in todo or motor.BIO_LINE not in todo:
        err.append('falta la línea de versión o la bio')
    if '1.0' not in motor.VERSION_FMT or 'octubre 2026' not in motor.VERSION_FMT:
        err.append('versión no es 1.0 · octubre 2026: %s' % motor.VERSION_FMT)
    if wf.sheetnames != [HM, HE, HI]:
        err.append('pestañas %r' % wf.sheetnames)
    if sum(1 for i in range(NI) if ING[i][0]) != NI or NI < 20:
        err.append('menos de 20 ingredientes')
    p = wf.properties
    if p.creator != 'AI Chef Pro' or p.description != 'aichef.pro/%s' % PID or 'AI Chef Pro' not in (p.keywords or ''):
        err.append('docProps: creator=%r description=%r keywords=%r' % (p.creator, p.description, p.keywords))
    for x in err:
        print('ERROR', x)
    print('autotest calculadora pizzería ES: %d comprobaciones (pycel + caché por defecto; pycel en 2 juegos más; B↔A) · %s'
          % (n, 'VERDE' if not err else 'ROJO (%d errores)' % len(err)))
    print('  masa por defecto: %d bolas × %g g → %.0f g · %.2f l · harina %.0f g · coste masa %.2f € · POR BOLA %.4f €'
          % (base['bolas'], base['peso'], r0['masa'], r0['litros'], r0['harina'], r0['coste_a'], r0['cbola_a']))
    fcs = [(pizzas[j], r0['p%d' % j]['fc'], r0['p%d' % j]['c_tot'], r0['p%d' % j]['sem']) for j in range(NP)]
    for nom, fc, ct, sem in fcs:
        print('  %-24s coste %.3f € · food cost %5.1f %% · %s' % (nom, ct, fc * 100, sem))
    return 1 if err else 0


if __name__ == '__main__':
    ap = argparse.ArgumentParser()
    ap.add_argument('--salida', required=True, help='carpeta de salida')
    ap.add_argument('--autotest', action='store_true')
    a = ap.parse_args()
    destino = os.path.join(a.salida, FICHERO)
    if not a.autotest and generar(destino):
        sys.exit('ERROR: fallos de pycel en inject_cache')
    sys.exit(autotest(destino))
