#!/usr/bin/env python3
"""
extraer_textos.py — HACCP Food Safety Kit Pro (EN) · cimientos de la F2-EN.

Copia adaptada de `restaurant-inventory-kit/extraer_textos.py` (sesión Claude Code, 3-oct-2026). Lee los
21 xlsx PUBLICADOS del Pack Plantillas APPCC v2.0 (`astro-site/public/dl/pack-appcc/`, SOLO LECTURA) y
escribe, junto a este script (o en --salida):

  · censo_es.json   por libro y hoja, lo que contarán los gates de la F2-EN: fórmulas (huella sha1 y
                    PATRONES normalizados con su recuento), referencias a hojas, literales en fórmulas y en
                    CF, DV (con `cita_temp` si el mensaje habla de °C), CF, gráficos, imágenes, merges,
                    protección, paneles, áreas y títulos de impresión, papel, cabeceras/pies, formatos,
                    fechas (celdas con formato de fecha y fechas escritas como TEXTO), números de ejemplo,
                    textos con °C o con normativa española, celdas verdes, columnas ocultas, docProps.
  · textos_es.json  cadenas de texto únicas con su contexto, TRATAMIENTO, PISTAS de la SPEC y GRUPO de
                    traducción (G1-G3 para los subagentes; GM = no se traduce).

    python3 extraer_textos.py [--origen DIR] [--salida DIR] [--md]

No traduce nada (regla 1bis: los textos EN los escriben subagentes Anthropic, nunca bridge.py).
Térmica: un solo proceso, en serie, sin builds ni navegador (~21 libros pequeños, < 10 s).

Tratamientos (la cadena toma el más «fuerte»: traducir > localizar > mapa-hoja > mapa-clave > regenerar):
  traducir    texto libre → subagente (G1-G3), con pistas de la SPEC.
  localizar   DATO DE EJEMPLO (proveedores, personas, platos, lotes, «(ejemplo)») → subagente del mismo
              grupo, con la tabla de datos de ejemplo de SPEC §5 (no traducción literal).
  mapa-hoja   nombre de pestaña o celda que lo cita → `mapas.HOJAS`.
  mapa-clave  ítem de DV o literal de fórmula/CF de ESA hoja (o celda de ejemplo igual a uno) → `mapas.CLAVES`.
  regenerar   lo escribe `aplicar_en.py`: versión, pie de página, docProps, línea de conservación
              (`mapas.CONSERVAR_EN`), tabla «Limits» del 02 y su lista en Instrucciones (`mapas.LIMITES_02`),
              gravedades del 15 (`mapas.SEVERIDAD_15`).
"""
import argparse
import datetime
import hashlib
import json
import os
import re
import subprocess
import sys
import zipfile
from collections import OrderedDict, defaultdict

import openpyxl
from openpyxl.styles.numbers import is_date_format
from openpyxl.utils import get_column_letter, range_boundaries

AQUI = os.path.dirname(os.path.abspath(__file__))
REPO = os.path.abspath(os.path.join(AQUI, '..', '..', '..'))
ORIGEN = os.path.join(REPO, 'astro-site', 'public', 'dl', 'pack-appcc')
sys.path.insert(0, AQUI)
import mapas                                                     # noqa: E402

VERDE = 'E8F5E9'
INVARIABLES = {'OK', 'N/A', 'AI Chef Pro', 'M', 'T', '☐'}
RX_VERSION = re.compile(r'^Versi[óo]n \d+\.\d+ · ')
RX_LITERAL = re.compile(r'"((?:[^"]|"")*)"')
RX_REF_HOJA = re.compile(r"(?:'((?:[^']|'')+)'|([^\W\d][\w.]*))!(?=\$?[A-Za-z]{1,3}\$?\d|\$?[A-Za-z]{1,3}:)")
RX_CITA = re.compile(r"«([^»]{1,40})»|\"([^\"]{1,40})\"")
RX_FECHA_TXT = re.compile(r'^\d{2}/\d{2}/\d{4}$')       # fecha ES escrita como TEXTO en la celda entera (D20)
RX_NORMA = re.compile(r'Reg\.|Rgto\.|Reglamento|\bRD\b|Real Decreto|Ley \d|Orden de|ROESB|\bLER\b|comunidad autónoma|'
                      r'Salud Pública|\b112\b|1169/2011|178/2002|852/2004|853/2004|2073/2005|Registro HA|'
                      r'registro sanitario|BOE')
RX_COMPARA = re.compile(r'[<>]=?\s*-?\d')

# Zonas de DATOS DE EJEMPLO por hoja: (fila1, fila2, columnas). Se localizan con SPEC §5.
ZONAS_EJEMPLO = {
    ('01', 'Registro Semanal'): (7, 13, 'DEFIJK'),
    ('02', 'Recepción Temperaturas'): (5, 7, 'ABCHIJ'),
    ('03', 'Productos químicos'): (5, 8, 'ABCDEGH'),
    ('04', 'Limpieza Diaria'): (38, 38, 'BG'),
    ('05', 'Recepción Mercancías'): (5, 7, 'ABCEJKLMN'),
    ('06', 'Trazabilidad'): (5, 7, 'ABCDEFGHI'),
    ('06', 'Salida y uso interno'): (5, 6, 'ABCDEFG'),
    ('07', 'Control Plagas DDD'): (5, 7, 'ACDEGHIJ'),
    ('07', 'Plano de cebos'): (5, 12, 'BDEGH'),
    ('08', 'Matriz Alérgenos'): (6, 13, 'BR'),
    ('09', 'Control Aceite'): (5, 7, 'ABGH'),
    ('09', 'Retirada de aceite usado'): (5, 5, 'ACDEF'),
    ('10', 'Control Agua'): (5, 7, 'ABG'),
    ('11', 'Acciones Correctivas'): (5, 11, 'BEFHIJKL'),
    ('16', 'Cocción y Regeneración'): (5, 7, 'BGIJ'),
    ('17', 'Enfriamiento'): (5, 7, 'BIJK'),
    ('17', 'Descongelación'): (5, 6, 'BCIJ'),
    ('18', 'Congelación Anisakis'): (5, 6, 'BCIJK'),
    ('19', 'Verificación Termómetros'): (5, 7, 'BHIJ'),
    ('B01', 'Formación Personal'): (5, 7, 'ABGHK'),
}
# Celdas que reescribe aplicar_en.py (no se traducen)
ZONAS_REGENERAR = {
    ('02', 'Límites'): (5, 14, 'AC', 'D10 tabla «Limits» (mapas.LIMITES_02)'),
    ('02', 'Instrucciones'): (13, 22, 'B', 'D10 lista de límites en °F (mapas.LIMITES_02)'),
    ('15', '25 Puntos Inspección'): (6, 34, 'D', 'D12 gravedad por punto (mapas.SEVERIDAD_15)'),
    ('02', 'Recepción Temperaturas'): (5, 7, 'E', 'D10 familia del ejemplo (mapas.POR_CELDA)'),
    ('16', 'Cocción y Regeneración'): (5, 7, 'C', 'D13 proceso del ejemplo (mapas.EJEMPLOS_NUM)'),
    ('17', 'Enfriamiento'): (4, 7, 'I', 'D14 columna I = lectura de 6 h (mapas.POR_CELDA, EJEMPLOS_NUM)'),
}

PISTAS = [
    (r'°C|Tª|T\.ª|[Tt]emperatura', 'D9: °F primero y °C entre paréntesis solo en texto («41 °F (5 °C)»); las cifras '
                                  'salen de SPEC §3 (FDA Food Code 2022), nunca de convertir el límite español'),
    (RX_NORMA.pattern, 'D8: la norma española/UE se sustituye por la referencia US (FDA Food Code 2022 §…, CFR, EPA) '
                       'con nota UK (FSA / SFBB) según SPEC §3; autoridad = «your local or state health department» '
                       '(UK: «your local authority»); 112 → 911 (UK 999)'),
    (r'APPCC|\bPCC\b|PPRo', 'Glosario: APPCC → HACCP; PCC → CCP; PPRo → OPRP (prerrequisito operativo)'),
    (r'[Aa]nisakis', 'D15: anisakis → «parasite destruction» (Food Code §3-402.11): −4 °F 168 h o −31 °F 15 h; '
                     'UK/UE 24 h a −20 °C'),
    (r'carn[ée]|manipulador', 'D18: no hay «carné» federal: CFPM (Food Code §2-102.12) + food handler card donde el '
                              'estado la exija; UK: sin certificado legal, formación «Level 2» habitual'),
    (r'48 h|48 horas', 'D18: vuelta al trabajo US 24 h sin síntomas (Food Code §2-201.13); UK 48 h (FSA)'),
    (r'A4', 'D21: papel US Letter'),
    (r'albar[áa]n', 'Glosario: albarán → «invoice / packing slip»'),
    (r'\bFDS\b|ficha de datos de seguridad', 'Glosario: FDS → SDS (Safety Data Sheet; UK: COSHH data sheet)'),
    (r'lej[íi]a|hipoclorito', 'Glosario: lejía/hipoclorito → «bleach (sodium hypochlorite)»'),
    (r'DDD|ROESB|biocida', 'D: empresa DDD con ROESB → «licensed pest control operator» (licencia estatal de '
                           'aplicador); nº de biocida → «EPA Reg. No.»; UK: BPCA/NPTA, autorización HSE'),
    (r'«|»', 'D6: pestañas y valores citados entre comillas dobles rectas ("Weekly Log")'),
    (r'registro \d\d|fichero \d\d|ficha \d\d|BONUS-0\d', 'D5: cita de plantilla → número + título EN (mapas.TITULOS)'),
    (r'\(ejemplo\)', '«(ejemplo)» → «(sample)»'),
    (r'\d,\d', 'Formato US: punto decimal y separador de miles con coma'),
    (r'\d{2}/\d{2}/\d{4}', 'D21: fechas de texto en formato US (mm/dd/yyyy)'),
    (r'€|euros?', 'Sin moneda (la tienda vende en USD; las plantillas no llevan importes)'),
    (r'Pack de Plantillas APPCC|Pack APPCC', 'D2: «HACCP Food Safety Kit Pro»'),
    (r'aichef\.pro/pack-appcc', 'D21: ' + mapas.URL_PRODUCTO),
    (r'Madrid', 'D17: ejemplo de altitud → Denver (5,280 ft): el agua hierve a unos 202 °F'),
    (r'[Aa]l[ée]rgeno', 'D8: US = 9 alérgenos mayores (FALCPA + FASTER Act, incl. sésamo) · UK = 14; la ficha y '
                        'la matriz conservan los 14 y marcan los 9 US'),
    (r'[Ee]nfriamiento|enfriar', 'D14: enfriamiento en 2 tramos: 135→70 °F en ≤ 2 h y →41 °F en ≤ 6 h en total '
                                 '(Food Code §3-501.14); UK: < 8 °C en 90 min (FSA)'),
    (r'[Cc]occi[óo]n|[Rr]egeneraci[óo]n|recalent', 'D13: cocción por tipo (165/155/145/135 °F, §3-401.11) y '
                                                   'recalentado a 165 °F en ≤ 2 h (§3-403.11)'),
    (r'[Aa]ceite|polares', 'D11: sin límite federal US de compuestos polares (TPM): 25 % se presenta como punto '
                           'de cambio del kit; máx. de fritura 375 °F'),
    (r'[Cc]loro|agua potable|pozo|dep[óo]sito', 'D: cloro libre 0,2-4,0 mg/L (EPA); pozo propio → análisis anual '
                                                '(Food Code §5-102.13)'),
    (r'[Ii]nspecci[óo]n|[Ii]nspector|grave', 'D12: gravedad = categoría Food Code (Priority / Priority foundation '
                                            '/ Core); UK: Food Hygiene Rating Scheme 0-5'),
    (r'[Tt]razabilidad', 'D: US FSMA 204 (registros al FDA en 24 h; obligatorio desde el 20-jul-2028 para los '
                         'alimentos de la Food Traceability List); UK: «one step back, one step forward»'),
]


# --------------------------------------------------------------------------
# utilidades
# --------------------------------------------------------------------------
def sha256(path):
    h = hashlib.sha256()
    with open(path, 'rb') as fh:
        for b in iter(lambda: fh.read(65536), b''):
            h.update(b)
    return h.hexdigest()


def corto(fichero):
    m = re.match(r'^(BONUS-)?(\d+)', fichero)
    return ('B' if m.group(1) else '') + m.group(2)


def tiene_letra(s):
    return any(ch.isalpha() for ch in s)


def literales(formula):
    return [m.group(1).replace('""', '"') for m in RX_LITERAL.finditer(formula or '')]


def refs_hoja(formula):
    sin_lit = RX_LITERAL.sub('""', formula or '')
    return [(m.group(1) or m.group(2)).replace("''", "'") for m in RX_REF_HOJA.finditer(sin_lit)]


def celdas_de(sqref):
    s = set()
    for rng in str(sqref).split():
        c1, r1, c2, r2 = range_boundaries(rng)
        for r in range(r1, r2 + 1):
            for c in range(c1, c2 + 1):
                s.add((r, c))
    return s


def en_zona(zonas, clave, cel):
    z = zonas.get(clave)
    if not z:
        return None
    if z[0] <= cel.row <= z[1] and get_column_letter(cel.column) in z[2]:
        return z
    return None


def claves_de_hoja(ws):
    """Ítems de DV de lista y literales de fórmula/CF de la hoja: lo que una celda puede «ser»."""
    s = set()
    for dv in ws.data_validations.dataValidation:
        f1 = dv.formula1 or ''
        if f1.startswith('"'):
            s.update(f1.strip('"').split(','))
    for row in ws.iter_rows():
        for c in row:
            if isinstance(c.value, str) and c.data_type == 'f':
                s.update(literales(c.value))
    for rng, reglas in ws.conditional_formatting._cf_rules.items():
        for rg in reglas:
            for f in rg.formula or []:
                s.update(literales(f))
    return s


def clasificar_celda(fichero, ws, cel, texto, hojas_libro, claves_hoja):
    f = corto(fichero)
    if RX_VERSION.match(texto):
        return 'regenerar', 'D21 línea de versión'
    if texto == mapas.PIE_ES:
        return 'regenerar', 'D2 pie «%s»' % mapas.PIE_EN
    if texto.replace('▸ ', '', 1) == mapas.CONSERVAR_ES:
        return 'regenerar', 'D19 línea de conservación (mapas.CONSERVAR_EN)'
    z = en_zona(ZONAS_REGENERAR, (f, ws.title), cel)
    if z:
        return 'regenerar', z[3]
    if texto in hojas_libro:
        return 'mapa-hoja', 'celda que cita la pestaña'
    if texto in claves_hoja and texto in mapas.CLAVES:
        return 'mapa-clave', 'valor de DV/literal de la hoja'
    if en_zona(ZONAS_EJEMPLO, (f, ws.title), cel) or '(ejemplo)' in texto:
        return 'localizar', 'dato de ejemplo (SPEC §5)'
    return 'traducir', None


# --------------------------------------------------------------------------
# censo de un libro
# --------------------------------------------------------------------------
def censar_libro(path, fichero):
    wb = openpyxl.load_workbook(path)
    hojas = list(wb.sheetnames)
    ocurr = []
    info = OrderedDict()
    info['corto'] = corto(fichero)
    info['sha256'] = sha256(path)
    info['nombre_en'] = mapas.FICHEROS[fichero]
    info['titulo_en'] = mapas.TITULOS[mapas.FICHEROS[fichero]]
    info['hojas'] = hojas
    info['hojas_en'] = [mapas.HOJAS.get(h) for h in hojas]
    info['nombres_definidos'] = sorted(getattr(wb.defined_names, 'keys', lambda: [])())
    pr = wb.properties
    info['docprops'] = OrderedDict((k, getattr(pr, k)) for k in
                                   ('creator', 'title', 'subject', 'keywords', 'description', 'category',
                                    'lastModifiedBy'))
    info['hojas_detalle'] = OrderedDict()
    for h in hojas:
        ocurr.append((h, {'f': corto(fichero), 'h': h, 't': 'hoja', 'tr': 'mapa-hoja', 'det': 'nombre de pestaña'}))

    fx_hoja_total = 0
    for ws in wb.worksheets:
        hd = OrderedDict()
        hd['nombre_en'] = mapas.HOJAS.get(ws.title)
        hd['dimension'] = ws.dimensions
        hd['max_fila'] = ws.max_row
        hd['max_col'] = ws.max_column
        claves_hoja = claves_de_hoja(ws)
        n_txt = n_num = n_fx = fx_con_hoja = 0
        refs = defaultdict(int)
        lits = defaultdict(list)
        formatos = defaultdict(int)
        fechas, fechas_txt, numeros, celsius, norma = [], [], [], [], []
        verdes = verdes_desbl = desbl = 0
        patrones = OrderedDict()
        huella = hashlib.sha1()
        for row in ws.iter_rows():
            for cel in row:
                if cel.__class__.__name__ == 'MergedCell':
                    continue
                v = cel.value
                fmt = cel.number_format
                fid = getattr(cel._style, 'numFmtId', None)
                if is_date_format(fmt):
                    fechas.append(OrderedDict([('celda', cel.coordinate), ('numFmtId', fid), ('formato', fmt),
                                               ('vacia', v is None)]))
                if fmt and fmt != 'General':
                    formatos[fmt] += 1
                fill = cel.fill.fgColor.rgb if cel.fill is not None and cel.fill.fgColor is not None else None
                bloq = cel.protection is None or cel.protection.locked is not False
                if not bloq:
                    desbl += 1
                if isinstance(fill, str) and fill.upper().endswith(VERDE):
                    verdes += 1
                    if not bloq:
                        verdes_desbl += 1
                if v is None:
                    continue
                if isinstance(v, str) and cel.data_type == 'f':
                    n_fx += 1
                    huella.update('{}={}\n'.format(cel.coordinate, v).encode('utf-8'))
                    for lit in literales(v):
                        lits[lit].append(cel.coordinate)
                    rh = refs_hoja(v)
                    if rh:
                        fx_con_hoja += 1
                    for nom in rh:
                        refs[nom] += 1
                    p = mapas.patron(v)
                    if p not in patrones:
                        patrones[p] = OrderedDict([('patron', p), ('n', 0), ('celdas', []),
                                                   ('cifras', bool(RX_COMPARA.search(RX_LITERAL.sub('""', v))))])
                    patrones[p]['n'] += 1
                    patrones[p]['celdas'].append(cel.coordinate)
                    continue
                if isinstance(v, str):
                    n_txt += 1
                    if '°C' in v or 'Tª' in v or 'T.ª' in v:
                        celsius.append(cel.coordinate)
                    if RX_NORMA.search(v):
                        norma.append(cel.coordinate)
                    if RX_FECHA_TXT.match(v.strip()):
                        fechas_txt.append(OrderedDict([('celda', cel.coordinate), ('texto', v)]))
                    tr, det = clasificar_celda(fichero, ws, cel, v, hojas, claves_hoja)
                    d = {'f': corto(fichero), 'h': ws.title, 'c': cel.coordinate, 't': 'celda', 'tr': tr}
                    if det:
                        d['det'] = det
                    ocurr.append((v, d))
                else:
                    n_num += 1
                    if ws.title != 'Instrucciones':
                        val = v.isoformat() if hasattr(v, 'isoformat') else v
                        numeros.append([cel.coordinate, val])
        fx_hoja_total += fx_con_hoja
        for p in patrones.values():
            cs = p['celdas']
            p['celdas'] = cs[0] + ('..' + cs[-1] if len(cs) > 1 else '')
        hd['celdas_texto'] = n_txt
        hd['celdas_numero'] = n_num
        hd['celdas_formula'] = n_fx
        hd['formulas_con_hoja'] = fx_con_hoja
        hd['refs_hoja'] = OrderedDict(sorted(refs.items()))
        hd['literales'] = OrderedDict((k, lits[k]) for k in sorted(lits))
        hd['literales_recuento'] = OrderedDict((k, len(lits[k])) for k in sorted(lits))
        hd['huella_formulas_sha1'] = huella.hexdigest() if n_fx else None
        hd['patrones_formula'] = list(patrones.values())
        hd['fechas'] = fechas
        hd['fechas_texto'] = fechas_txt
        hd['numeros'] = numeros
        hd['formatos'] = OrderedDict(sorted(formatos.items(), key=lambda kv: -kv[1]))
        hd['textos_celsius'] = celsius
        hd['textos_normativa_es'] = norma
        hd['verdes'] = verdes
        hd['verdes_desbloqueadas'] = verdes_desbl
        hd['desbloqueadas'] = desbl
        hd['columnas_ocultas'] = sorted(k for k, d in ws.column_dimensions.items() if d.hidden)
        hd['filas_ocultas'] = sorted(k for k, d in ws.row_dimensions.items() if d.hidden)
        hd['merges'] = sorted(str(m) for m in ws.merged_cells.ranges)
        hd['paneles'] = ws.freeze_panes
        p = ws.protection
        hd['proteccion'] = OrderedDict([('hoja', bool(p.sheet)), ('password', bool(p.password))])
        hd['autofiltro'] = ws.auto_filter.ref
        hd['area_impresion'] = ws.print_area
        hd['titulos_impresion'] = OrderedDict([('filas', ws.print_title_rows), ('columnas', ws.print_title_cols)])
        ps = ws.page_setup
        hd['papel'] = OrderedDict([('paperSize', ps.paperSize), ('orientation', ps.orientation),
                                   ('fitToWidth', ps.fitToWidth), ('fitToHeight', ps.fitToHeight),
                                   ('fitToPage', bool(ws.sheet_properties.pageSetUpPr and
                                                      ws.sheet_properties.pageSetUpPr.fitToPage))])
        hf = OrderedDict()
        for nombre in ('oddHeader', 'oddFooter', 'evenHeader', 'evenFooter', 'firstHeader', 'firstFooter'):
            obj = getattr(ws, nombre)
            for parte in ('left', 'center', 'right'):
                txt = getattr(obj, parte).text
                if txt:
                    hf['{}.{}'.format(nombre, parte)] = txt
                    ocurr.append((txt, {'f': corto(fichero), 'h': ws.title, 't': 'cabecera-pie',
                                        'c': '{}.{}'.format(nombre, parte), 'tr': 'regenerar',
                                        'det': 'D21 pie «%s»' % mapas.FOOTER_EN}))
        hd['cabeceras_pies'] = hf
        dvs = []
        for dv in ws.data_validations.dataValidation:
            msg = ' '.join(x for x in (dv.error, dv.prompt, dv.errorTitle) if x)
            dvs.append(OrderedDict([
                ('sqref', str(dv.sqref)), ('type', dv.type), ('operator', dv.operator),
                ('formula1', dv.formula1), ('formula2', dv.formula2), ('allow_blank', dv.allow_blank),
                ('showErrorMessage', dv.showErrorMessage), ('errorStyle', dv.errorStyle),
                ('errorTitle', dv.errorTitle), ('error', dv.error),
                ('promptTitle', dv.promptTitle), ('prompt', dv.prompt),
                ('n_celdas', len(celdas_de(dv.sqref))),
                ('cita_hojas', sorted(set(refs_hoja(dv.formula1 or '')))),
                ('cita_temp', '°C' in msg),
            ]))
            base = {'f': corto(fichero), 'h': ws.title, 'c': str(dv.sqref)}
            for campo, t in (('errorTitle', 'dv-error-titulo'), ('error', 'dv-error'),
                             ('promptTitle', 'dv-prompt-titulo'), ('prompt', 'dv-prompt')):
                txt = getattr(dv, campo)
                if txt:
                    ocurr.append((txt, dict(base, t=t, tr='traducir')))
            f1 = dv.formula1 or ''
            if f1.startswith('"'):
                for it in f1.strip('"').split(','):
                    if it in mapas.CLAVES:
                        d = dict(base, t='dv-item', tr='mapa-clave', det='mapas.CLAVES')
                    else:
                        d = dict(base, t='dv-item', tr='regenerar', det='D13 lista especial (mapas.DV_LISTAS_ESPECIALES)')
                    ocurr.append((it, d))
        hd['dv'] = dvs
        cfs = []
        lits_cf = defaultdict(int)
        for rng, reglas in ws.conditional_formatting._cf_rules.items():
            for rg in reglas:
                fs = list(rg.formula or [])
                cfs.append(OrderedDict([
                    ('sqref', str(rng.sqref)), ('type', rg.type), ('operator', rg.operator),
                    ('text', getattr(rg, 'text', None)), ('formula', fs), ('priority', rg.priority),
                    ('stopIfTrue', rg.stopIfTrue), ('cita_hojas', sorted({h for f in fs for h in refs_hoja(f)})),
                ]))
                for f in fs:
                    for lit in literales(f):
                        lits_cf[lit] += 1
        hd['cf'] = cfs
        hd['literales_cf'] = OrderedDict(sorted(lits_cf.items()))
        info['hojas_detalle'][ws.title] = hd

    with zipfile.ZipFile(path) as z:
        nombres = z.namelist()
        info['graficos'] = sorted(n for n in nombres if re_chart(n))
        info['imagenes'] = sorted(n for n in nombres if n.startswith('xl/media/'))
    for campo, v in info['docprops'].items():
        if not isinstance(v, str) or not v:
            continue
        tr, det = 'regenerar', 'D21 docProps'
        if campo == 'keywords':
            tr, det = 'traducir', None
        if campo == 'creator':
            continue
        d = {'f': corto(fichero), 't': 'docprops', 'c': campo, 'tr': tr}
        if det:
            d['det'] = det
        ocurr.append((v, d))

    hdv = list(info['hojas_detalle'].values())
    lit_tot, litcf_tot = defaultdict(int), defaultdict(int)
    for x in hdv:
        for k, v in x['literales_recuento'].items():
            lit_tot[k] += v
        for k, v in x['literales_cf'].items():
            litcf_tot[k] += v
    info['totales'] = OrderedDict([
        ('hojas', len(hojas)),
        ('celdas_texto', sum(x['celdas_texto'] for x in hdv)),
        ('celdas_formula', sum(x['celdas_formula'] for x in hdv)),
        ('patrones_formula', sum(len(x['patrones_formula']) for x in hdv)),
        ('formulas_con_hoja', fx_hoja_total),
        ('dv', sum(len(x['dv']) for x in hdv)),
        ('dv_lista_literal', sum(1 for x in hdv for d in x['dv'] if (d['formula1'] or '').startswith('"'))),
        ('dv_cita_hoja', sum(1 for x in hdv for d in x['dv'] if d['cita_hojas'])),
        ('dv_temperatura', sum(1 for x in hdv for d in x['dv'] if d['cita_temp'])),
        ('dv_celdas', sum(d['n_celdas'] for x in hdv for d in x['dv'])),
        ('cf', sum(len(x['cf']) for x in hdv)),
        ('graficos', len(info['graficos'])),
        ('imagenes', len(info['imagenes'])),
        ('merges', sum(len(x['merges']) for x in hdv)),
        ('hojas_protegidas', sum(1 for x in hdv if x['proteccion']['hoja'])),
        ('areas_impresion', sum(1 for x in hdv if x['area_impresion'])),
        ('titulos_impresion', sum(1 for x in hdv if x['titulos_impresion']['filas'])),
        ('paneles', sum(1 for x in hdv if x['paneles'])),
        ('fechas', sum(len(x['fechas']) for x in hdv)),
        ('fechas_texto', sum(len(x['fechas_texto']) for x in hdv)),
        ('numeros_ejemplo', sum(len(x['numeros']) for x in hdv)),
        ('textos_celsius', sum(len(x['textos_celsius']) for x in hdv)),
        ('textos_normativa_es', sum(len(x['textos_normativa_es']) for x in hdv)),
        ('verdes', sum(x['verdes'] for x in hdv)),
        ('verdes_desbloqueadas', sum(x['verdes_desbloqueadas'] for x in hdv)),
        ('desbloqueadas', sum(x['desbloqueadas'] for x in hdv)),
        ('columnas_ocultas', sum(len(x['columnas_ocultas']) for x in hdv)),
        ('papel', sorted({x['papel']['paperSize'] for x in hdv if x['papel']['paperSize'] is not None})),
        ('literales', OrderedDict(sorted(lit_tot.items()))),
        ('literales_cf', OrderedDict(sorted(litcf_tot.items()))),
    ])
    return info, ocurr


def re_chart(n):
    return re.fullmatch(r'xl/charts/chart\d+\.xml', n) is not None


# --------------------------------------------------------------------------
# cadenas, grupos y pistas
# --------------------------------------------------------------------------
PRIORIDAD = {'traducir': 0, 'localizar': 1, 'mapa-hoja': 2, 'mapa-clave': 3, 'regenerar': 4}
SUMABLES = ('hojas', 'celdas_texto', 'celdas_formula', 'patrones_formula', 'formulas_con_hoja', 'dv',
            'dv_lista_literal', 'dv_cita_hoja', 'dv_temperatura', 'dv_celdas', 'cf', 'graficos', 'imagenes',
            'merges', 'hojas_protegidas', 'areas_impresion', 'titulos_impresion', 'paneles', 'fechas',
            'fechas_texto', 'numeros_ejemplo', 'textos_celsius', 'textos_normativa_es', 'verdes',
            'verdes_desbloqueadas', 'desbloqueadas', 'columnas_ocultas')


def pistas(texto):
    return [nota for rx, nota in PISTAS if re.search(rx, texto)]


def citas_hoja(texto, todas_hojas):
    out = []
    for m in RX_CITA.finditer(texto):
        span = next(g for g in m.groups() if g is not None)
        if span in todas_hojas and span not in [c['es'] for c in out]:
            out.append(OrderedDict([('es', span), ('en', mapas.HOJAS.get(span))]))
    return out


def palabras(s):
    return len(re.findall(r'\w+', s))


def tres_bloques(orden, peso):
    """Parte la lista de libros en 3 bloques contiguos minimizando el bloque más pesado."""
    n = len(orden)
    mejor = None
    for i in range(1, n - 1):
        for j in range(i + 1, n):
            bl = [orden[:i], orden[i:j], orden[j:]]
            m = max(sum(peso[f] for f in b) + (peso['G1-comun'] if k == 0 else 0) for k, b in enumerate(bl))
            if mejor is None or m < mejor[0]:
                mejor = (m, bl)
    return mejor[1]


def main():
    ap = argparse.ArgumentParser(description=__doc__.split('\n')[1])
    ap.add_argument('--origen', default=ORIGEN)
    ap.add_argument('--salida', default=AQUI)
    ap.add_argument('--md', action='store_true', help='imprime la tabla §0 de F1-inventario-es.md')
    args = ap.parse_args()

    libros = list(mapas.LIBROS_ES)
    faltan = [f for f in libros if not os.path.exists(os.path.join(args.origen, f))]
    if faltan:
        sys.exit('faltan ficheros en {}: {}'.format(args.origen, faltan))
    sobran = sorted(f for f in os.listdir(args.origen) if f.endswith('.xlsx') and f not in libros)
    if sobran:
        sys.exit('xlsx no previstos en {}: {}'.format(args.origen, sobran))
    try:
        commit = subprocess.check_output(['git', '-C', REPO, 'rev-parse', '--short', 'HEAD'],
                                         stderr=subprocess.DEVNULL).decode().strip()
    except Exception:                                             # noqa: BLE001
        commit = None

    censo = OrderedDict()
    censo['meta'] = OrderedDict([
        ('generado', datetime.datetime.utcnow().replace(microsecond=0).isoformat() + 'Z'),
        ('script', 'scripts/productos-digitales/haccp-kit/extraer_textos.py'),
        ('origen', os.path.relpath(args.origen, REPO)),
        ('commit', commit),
        ('nota', 'Censo del Pack Plantillas APPCC v2.0 PUBLICADO. Los gates de la F2-EN comparan contra este '
                 'fichero, nunca contra recuentos fijos. patrones_formula = fórmulas con las filas normalizadas '
                 'a {r} (mapas.patron); «cifras» = el patrón compara con un número (lo decide mapas.FORMULAS_EN '
                 'o mapas.PATRONES_SIN_CAMBIO).'),
    ])
    censo['libros'] = OrderedDict()
    todas = []
    for f in libros:
        info, oc = censar_libro(os.path.join(args.origen, f), f)
        censo['libros'][f] = info
        todas.extend(oc)
    T = OrderedDict((k, sum(censo['libros'][f]['totales'][k] for f in libros)) for k in SUMABLES)
    lit, litcf = defaultdict(int), defaultdict(int)
    for f in libros:
        for k, v in censo['libros'][f]['totales']['literales'].items():
            lit[k] += v
        for k, v in censo['libros'][f]['totales']['literales_cf'].items():
            litcf[k] += v
    T['literales'] = OrderedDict(sorted(lit.items(), key=lambda kv: -kv[1]))
    T['literales_cf'] = OrderedDict(sorted(litcf.items(), key=lambda kv: -kv[1]))
    T['papel'] = sorted({p for f in libros for p in censo['libros'][f]['totales']['papel']})
    censo['totales'] = T

    # ---- cadenas únicas
    todas_hojas = set(h for f in libros for h in censo['libros'][f]['hojas'])
    excl = OrderedDict()
    cad = OrderedDict()
    for texto, oc in todas:
        if not tiene_letra(texto) or texto in INVARIABLES:
            excl[texto] = excl.get(texto, 0) + 1
            continue
        cad.setdefault(texto, []).append(oc)
    orden_f = [corto(f) for f in libros]
    cadenas = []
    for i, (texto, ocs) in enumerate(cad.items(), 1):
        fich = []
        for o in ocs:
            if o['f'] not in fich:
                fich.append(o['f'])
        trs = sorted({o['tr'] for o in ocs}, key=lambda t: PRIORIDAD[t])
        e = OrderedDict()
        e['id'] = 'c{:04d}'.format(i)
        e['es'] = texto
        e['grupo'] = None
        e['tratamiento'] = trs[0]
        if len(trs) > 1:
            e['tratamientos'] = trs
        if texto in mapas.HOJAS and 'mapa-hoja' in trs:
            e['en_mapa'] = mapas.HOJAS[texto]
        elif texto in mapas.CLAVES and 'mapa-clave' in trs:
            e['en_mapa'] = mapas.CLAVES[texto]
        e['n_ficheros'] = len(fich)
        e['n_apariciones'] = len(ocs)
        e['palabras'] = palabras(texto)
        e['con_cifra'] = bool(re.search(r'\d', texto))
        ch = citas_hoja(texto, todas_hojas)
        if ch:
            e['cita_hojas'] = ch
        if trs[0] in ('traducir', 'localizar'):
            pi = pistas(texto)
            if pi:
                e['pistas'] = pi
        e['sha1'] = hashlib.sha1(texto.encode('utf-8')).hexdigest()[:12]
        e['donde'] = ocs
        e['_ficheros'] = fich
        cadenas.append(e)

    # ---- grupos: G1-G3 (subagentes) por bloques contiguos de libros; comunes → G1; GM no se traduce
    peso = defaultdict(int)
    for e in cadenas:
        if e['tratamiento'] in ('traducir', 'localizar'):
            comun = e['n_ficheros'] >= 3
            e['_comun'] = comun
            peso['G1-comun' if comun else e['_ficheros'][0]] += e['palabras']
    bloques = tres_bloques(orden_f, peso)
    gid_de = {}
    for k, bl in enumerate(bloques, 1):
        for f in bl:
            gid_de[f] = 'G%d' % k
    for e in cadenas:
        if e['tratamiento'] in ('traducir', 'localizar'):
            e['grupo'] = 'G1' if e['_comun'] else gid_de[e['_ficheros'][0]]
            if e['_comun']:
                e['comun'] = True
        else:
            e['grupo'] = 'GM'
    grupos = []
    for k, bl in enumerate(bloques, 1):
        gid = 'G%d' % k
        es = [e for e in cadenas if e['grupo'] == gid]
        grupos.append(OrderedDict([
            ('id', gid),
            ('descripcion', ('Cadenas comunes a ≥ 3 libros (se traducen PRIMERO: fijan el glosario) + ' if k == 1 else '')
             + 'libros ' + ', '.join(bl)),
            ('libros', bl),
            ('n_cadenas', len(es)), ('n_localizar', sum(1 for e in es if e['tratamiento'] == 'localizar')),
            ('n_palabras', sum(e['palabras'] for e in es)),
        ]))
    gm = [e for e in cadenas if e['grupo'] == 'GM']
    por_tr = defaultdict(int)
    for e in gm:
        por_tr[e['tratamiento']] += 1
    grupos.append(OrderedDict([
        ('id', 'GM'),
        ('descripcion', 'NO se traduce: pestañas y claves (mapas.py), versión, pie, docProps, línea de conservación, '
                        'tabla Limits del 02, gravedades del 15. ' +
                        ', '.join('{} {}'.format(v, k) for k, v in sorted(por_tr.items()))),
        ('n_cadenas', len(gm)), ('n_palabras', sum(e['palabras'] for e in gm)),
    ]))
    orden_g = {g['id']: k for k, g in enumerate(grupos)}
    cadenas.sort(key=lambda e: (orden_g[e['grupo']], int(e['id'][1:])))
    for e in cadenas:
        e.pop('_ficheros')
        e.pop('_comun', None)
    assert all(e['grupo'] for e in cadenas)

    textos = OrderedDict()
    textos['meta'] = OrderedDict([
        ('generado', censo['meta']['generado']),
        ('script', censo['meta']['script']),
        ('origen', censo['meta']['origen']),
        ('commit', commit),
        ('libros', OrderedDict((corto(f), OrderedDict([('es', f), ('en', mapas.FICHEROS[f]),
                                                      ('sha256', censo['libros'][f]['sha256'])]))
                               for f in libros)),
        ('n_cadenas', len(cadenas)),
        ('n_traducir', sum(1 for e in cadenas if e['tratamiento'] == 'traducir')),
        ('n_localizar', sum(1 for e in cadenas if e['tratamiento'] == 'localizar')),
        ('n_palabras_subagentes', sum(e['palabras'] for e in cadenas if e['grupo'] != 'GM')),
        ('n_apariciones', sum(e['n_apariciones'] for e in cadenas)),
        ('grupos', grupos),
        ('excluidas', OrderedDict([('cadenas', excl),
                                   ('regla', 'Fuera: cadenas sin letras y las invariables %s. Los literales de '
                                             'fórmula están en censo_es.json.' % sorted(INVARIABLES))])),
        ('campos', OrderedDict([
            ('donde', 'f = libro, h = hoja, c = celda / sqref / campo; t = celda | hoja | dv-* | cabecera-pie | '
                      'docprops; tr = tratamiento de esa aparición; det = detalle'),
            ('en_mapa', 'valor EN obligatorio que dicta mapas.py'),
            ('cita_hojas', 'pestañas ES citadas entre comillas, con su EN'),
            ('pistas', 'decisiones de la SPEC que aplican a la cadena'),
        ])),
    ])
    textos['cadenas'] = cadenas

    os.makedirs(args.salida, exist_ok=True)
    for nombre, obj in (('textos_es.json', textos), ('censo_es.json', censo)):
        with open(os.path.join(args.salida, nombre), 'w', encoding='utf-8') as fh:
            json.dump(obj, fh, ensure_ascii=False, indent=1)
            fh.write('\n')

    print('textos_es.json: {} cadenas ({} traducir, {} localizar, {} palabras para subagentes) · {} apariciones'.format(
        len(cadenas), textos['meta']['n_traducir'], textos['meta']['n_localizar'],
        textos['meta']['n_palabras_subagentes'], textos['meta']['n_apariciones']))
    for g in grupos:
        print('  {:<3} {:>4} cadenas {:>6} palabras  {}'.format(g['id'], g['n_cadenas'], g['n_palabras'],
                                                              g['descripcion'][:110]))
    print('censo_es.json: ' + json.dumps({k: v for k, v in T.items() if not isinstance(v, dict)},
                                         ensure_ascii=False))
    if args.md:
        print(tabla_md(censo))


COLS_MD = [('hojas', 'Hojas'), ('celdas_texto', 'Texto'), ('celdas_formula', 'Fórmulas'),
           ('formulas_con_hoja', 'Fx→hoja'), ('dv', 'DV'), ('cf', 'CF'), ('merges', 'Merges'),
           ('fechas', 'Fechas'), ('fechas_texto', 'Fechas texto'), ('numeros_ejemplo', 'Nº ejemplo'),
           ('textos_celsius', 'Textos °C'), ('textos_normativa_es', 'Norma ES'), ('verdes', 'Verdes'),
           ('areas_impresion', 'Áreas')]


def tabla_md(censo):
    out = ['| Libro | ' + ' | '.join(t for _, t in COLS_MD) + ' |', '|---|' + '---|' * len(COLS_MD)]
    for f, info in censo['libros'].items():
        out.append('| %s | ' % f[:-5] + ' | '.join(str(info['totales'][k]) for k, _ in COLS_MD) + ' |')
    out.append('| **Total** | ' + ' | '.join('**%s**' % censo['totales'][k] for k, _ in COLS_MD) + ' |')
    return '\n'.join(out)


if __name__ == '__main__':
    main()
