#!/usr/bin/env python3
"""
resumen_entregables.py — Bar Business Plan Kit (EN) · cifras y recuentos LEÍDOS de los 3 ficheros EN generados
(`astro-site/public/dl/bar-business-plan/`), para la F3 (landing, changelog, correo) y la revisión final. Sesión
Claude Code, 11-oct-2026 (cierre de la F2). Ejecutar en el VPS (openpyxl + python-docx):

    /root/venv-guias/bin/python resumen_entregables.py

Escribe `entregables_bar.json` al lado: por fichero, sha256, hojas, fórmulas, tareas por fase, párrafos y palabras del
docx; el caso (de `cifras_caso.json`, que es la caché del libro EN calibrado: texto US de cada token) y un barrido
independiente de cifras ES prohibidas (D52), restos de español y caracteres no latinos sobre TODO el texto de los 3.
"""
import hashlib
import json
import os
import re
import sys
from collections import OrderedDict

AQUI = os.path.dirname(os.path.abspath(__file__))
os.environ['BP_CONJUNTO'] = 'bar'
sys.path.insert(0, os.path.join(os.path.dirname(AQUI), 'business-plans-en'))
import mapas                                                     # noqa: E402
import openpyxl                                                  # noqa: E402
import docx                                                      # noqa: E402

RX_LIT = re.compile(r'"((?:[^"]|"")*)"')
REPO = os.path.abspath(os.path.join(AQUI, '..', '..', '..'))
DL = os.path.join(REPO, mapas.PLANES['bar']['dir_en'])
CASO = ('ventas_a1', 'ventas_a2', 'ventas_a3', 'clientes_dia', 'ticket', 'precio_menu', 'dias', 'bebida_pct',
        'alcohol_pct', 'aforo', 'rotaciones_dia', 'inversion_total', 'capex', 'fondo_maniobra', 'necesidad_caja',
        'inv_licencia', 'fondos_propios', 'prestamo', 'aportacion_pct', 'tipo_prestamo', 'plazo_prestamo',
        'cuota_anual', 'margen_bruto_pct', 'cogs_pct', 'personal_pct', 'prime_cost_pct', 'ocupacion_pct',
        'ebitda_a1', 'ebitda_pct_a1', 'resultado_neto_a1', 'resultado_neto_a2', 'resultado_neto_a3', 'margen_neto_a1',
        'dscr_a1', 'dscr_min', 'equilibrio_clientes_dia', 'equilibrio_caja_clientes_dia', 'holgura_caja_pct',
        'saldo_minimo', 'mes_saldo_minimo', 'payback', 'payback_capex', 'plantilla_puestos', 'plantilla_jornadas',
        'cobertura_horas', 'pesimista_ventas', 'pesimista_resultado', 'pesimista_caja', 'optimista_ventas',
        'optimista_resultado', 'seguros_anual')


def sha(p):
    return hashlib.sha256(open(p, 'rb').read()).hexdigest()


def barrido(textos):
    out = OrderedDict([('cifras_prohibidas', []), ('espanol', []), ('no_latinos', [])])
    for lugar, t in textos:
        for motivo, frag in mapas.cifras_prohibidas(t):
            out['cifras_prohibidas'].append('%s: %s «%s»' % (lugar, motivo, frag))
        for motivo, ctx in mapas.restos_espanol(t):
            if not motivo.startswith('oficio'):
                out['espanol'].append('%s: %s · %s' % (lugar, motivo, ctx.strip()[:50]))
        if mapas.no_latinos(t):
            out['no_latinos'].append('%s: %s' % (lugar, mapas.no_latinos(t)))
    return out


def main():
    res = OrderedDict([('producto', mapas.PLANES['bar']['producto']), ('carpeta', mapas.PLANES['bar']['dir_en']),
                       ('ficheros', OrderedDict())])
    textos = []
    for corto, (_plan, _fes, fen, tit) in mapas.LIBROS.items():
        p = os.path.join(DL, fen)
        wb = openpyxl.load_workbook(p)
        n_f = sum(1 for ws in wb.worksheets for c in ws._cells.values()
                  if isinstance(c.value, str) and c.value.startswith('='))
        d = OrderedDict([('titulo', tit), ('sha256', sha(p)), ('bytes', os.path.getsize(p)),
                         ('hojas', wb.sheetnames), ('n_hojas', len(wb.sheetnames)), ('formulas', n_f)])
        if mapas.TIPO_DE[corto] == 'check':
            tareas = mapas.contar_tareas(corto, wb)
            d['tareas_por_fase'] = tareas
            d['tareas'] = sum(tareas)
        res['ficheros'][fen] = d
        for ws in wb.worksheets:
            for c in ws._cells.values():
                if not isinstance(c.value, str):
                    continue
                lugar = '%s %s!%s' % (corto, ws.title, c.coordinate)
                if c.value.startswith('='):          # fórmula: solo sus literales (la coma separa argumentos)
                    textos += [(lugar + ' (literal)', x.replace('""', '"')) for x in RX_LIT.findall(c.value)]
                else:
                    textos.append((lugar, c.value))
    fen = mapas.DOCX['bar'][1]
    p = os.path.join(DL, fen)
    dd = docx.Document(p)
    pars = dd.paragraphs
    res['ficheros'][fen] = OrderedDict([
        ('titulo', mapas.DOCX['bar'][2]), ('sha256', sha(p)), ('bytes', os.path.getsize(p)),
        ('parrafos', len(pars)), ('con_texto', sum(1 for x in pars if x.text.strip())),
        ('palabras', sum(len(x.text.split()) for x in pars)),
        ('encabezados', [x.text for x in pars if x.style.name == 'Heading 1'])])
    textos += [('docx[%d]' % i, x.text) for i, x in enumerate(pars) if x.text.strip()]
    cif = json.load(open(os.path.join(AQUI, 'cifras_caso.json'), encoding='utf-8'))['bar']
    res['caso'] = OrderedDict((t, cif[t]['texto']) for t in CASO if t in cif)
    res['barrido'] = barrido(textos)
    res['barrido']['piezas'] = len(textos)
    with open(os.path.join(AQUI, 'entregables_bar.json'), 'w', encoding='utf-8') as fh:
        json.dump(res, fh, ensure_ascii=False, indent=1)
        fh.write('\n')
    b = res['barrido']
    print('entregables_bar.json · %d piezas · prohibidas %d · español %d · no latinos %d' % (
        b['piezas'], len(b['cifras_prohibidas']), len(b['espanol']), len(b['no_latinos'])))
    for k in ('cifras_prohibidas', 'espanol', 'no_latinos'):
        for x in b[k][:10]:
            print('  ', k, x)
    return 1 if (b['cifras_prohibidas'] or b['espanol'] or b['no_latinos']) else 0


if __name__ == '__main__':
    sys.exit(main())
