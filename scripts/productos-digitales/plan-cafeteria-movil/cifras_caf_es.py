#!/usr/bin/env python3
"""
cifras_caf_es.py — Cafetería Móvil para Eventos ES 1.0 (copia de `plan-pizzeria/cifras_piz_es.py` —redondeo comercial
con Decimal, m2— con el MAPA de eventos de `plan-finca-eventos/cifras_fin_es.py`; rótulos de línea A medidos sobre el
libro generado el 28-oct-2026, ya pasados por el VOCABULARIO de `contenido_plan_negocio_cafeteria_movil_eventos/a.py`):
lee la CACHÉ del plan financiero (después de `inject_cache`, que hace `main.py`) por RÓTULO, escribe
`cifras_caso_es.json` (los tokens {ASI} del Word) y pasa los gates del caso (SPEC D9 y F1-camino §6):
  · DSCR mínimo ≥ 1,25 · saldo mínimo de tesorería > 0 · cobertura de horas ≥ 98 %;
  · holgura sobre el equilibrio de caja del año 1 ≥ 15 % (las de los años 2 y 3, informativas, ≥ 0);
  · exceso de financiación (origen − usos) ≥ 0 y < 5 % (el mismo control que la demo RD-34 del motor);
  · resultado neto a régimen (años 2 y 3) entre el 7 y el 12 % de las ventas, y el del año 1 ≤ 12 %;
  · venta directa en ferias y mercados (`pct_delivery`) < 25 % (demo §2.11.9) y = la cuota de la mezcla ± 1 punto;
  · las dos líneas de venta al 10 % (columna G del P&L) y el IVA repercutido de restauración = 10 % (SPEC D6).
Con `--calc` lee también la caché de la calculadora (`calculadora_caf_es.py`, otro agente: hojas «Escandallo» y
«Ejemplos»): COSTE_ESPRESSO, COSTE_LATTE, COSTE_MEDIO_BEBIDA… y los tres presupuestos de ejemplo (no entran en los
gates). Corre en el VPS (openpyxl).

    /root/venv-guias/bin/python cifras_caf_es.py --xlsx <plan-financiero-cafeteria-movil-eventos.xlsx> \
        [--calc <calculadora-escandallo-presupuesto-cafe-eventos.xlsx>] [--salida <json>]

Sesión Claude Code, 28-oct-2026.
"""
import argparse
import json
import math
import os
import re
import sys
from decimal import ROUND_HALF_UP, Decimal

import openpyxl

AQUI = os.path.dirname(os.path.abspath(__file__))
SALIDA = os.path.join(AQUI, 'cifras_caso_es.json')

#: Mes 1 = marzo (SPEC D7): nombre del mes del calendario de cada «Mes n» de la Tesorería.
MESES = ('marzo', 'abril', 'mayo', 'junio', 'julio', 'agosto', 'septiembre', 'octubre', 'noviembre', 'diciembre',
         'enero', 'febrero')

#: Mezcla del caso (research §4 con la palanca 1 de D9: bodas a 9,00 €): (tipología, eventos, invitados, € sin IVA).
#: La venta directa es la última fila. Sirve para el gate de la cuota de ferias y para tokens del Word.
MEZCLA = (
    ('bodas', 30, 120, 9.00),
    ('empresa', 40, 80, 8.00),
    ('congresos', 22, 300, 4.50),
    ('privados', 25, 50, 9.50),
    ('venta_directa', 12, 250, 3.00),
)

#: token → (hoja, regex del rótulo de la columna A, columna, formato). Formatos: eur (miles con punto, sin
#: decimales) · eur2 · pct (1 decimal) · pct0 · dec1 · dec2 · int · int_up (redondeo hacia arriba).
SUP, PYG, INV, PE, TES, FIN, PER, ESC = ('0. Supuestos', 'PyG 3 Años', 'Inversión Inicial', 'Punto Equilibrio',
                                         'Tesorería 12 meses', 'Financiación', 'Personal', 'Escenarios')
MAPA = {
    # Supuestos (driver sin permutar: invitados por evento × precio por invitado × eventos al año)
    'INVITADOS_EVENTO': (SUP, r'^invitados por evento \(media', 2, 'int'),
    'PRECIO_INVITADO': (SUP, r'^precio medio por invitado sin iva', 2, 'eur2'),
    'PRECIO_INVITADO_IVA': (SUP, r'^pvp equivalente con iva', 2, 'eur2'),
    'EVENTOS_ANO': (SUP, r'^eventos al a[ñn]o$', 2, 'int'),
    'CREC_A2': (SUP, r'^crecimiento de volumen del a[ñn]o 2', 2, 'pct0'),
    'CREC_A3': (SUP, r'^crecimiento de volumen del a[ñn]o 3', 2, 'pct0'),
    'PCT_DULCES': (SUP, r'^ventas de comida sobre el total', 2, 'pct0'),
    'PCT_BEBIDAS': (SUP, r'^ventas de bebida sobre el total', 2, 'pct0'),
    'COSTE_DULCES': (SUP, r'^coste de mercanc[ií]a sobre las ventas de comida', 2, 'pct0'),
    'COSTE_BEBIDAS': (SUP, r'^coste de mercanc[ií]a sobre las ventas de bebida', 2, 'pct0'),
    'CONSUMIBLES_PCT': (SUP, r'^consumibles sobre ventas', 2, 'pct0'),
    'CUOTA_FERIAS': (SUP, r'^ventas directas en ferias y mercados', 2, 'pct'),
    'CANON_FERIAS': (SUP, r'^canon o comisi[oó]n del organizador', 2, 'pct0'),
    'COMISION_TPV': (SUP, r'^comisi[oó]n de los medios de pago', 2, 'pct'),
    'CUOTA_BASE_MES': (SUP, r'^cuota mensual del almac[eé]n base', 2, 'eur'),
    'FIANZA_MESES': (SUP, r'^fianza del almac[eé]n base \(meses', 2, 'int'),
    'SUMINISTROS_MES': (SUP, r'^suministros mensuales', 2, 'eur'),
    'SEGUROS_ANO': (SUP, r'^seguros', 2, 'eur'),
    'VARIOS_PCT': (SUP, r'^varios e imprevistos sobre ventas', 2, 'pct0'),
    'RECURSOS_PROPIOS': (SUP, r'^recursos propios', 2, 'eur'),
    'TIPO_PRESTAMO': (SUP, r'^tipo de inter[eé]s nominal', 2, 'pct'),
    'PLAZO_PRESTAMO': (SUP, r'^plazo del pr[eé]stamo', 2, 'int'),
    'CARENCIA': (SUP, r'^carencia de principal', 2, 'int'),
    'MESES_FONDO': (SUP, r'^fondo de maniobra \(meses', 2, 'int'),
    'IS_NUEVA': (SUP, r'^impuesto de sociedades, nueva creaci', 2, 'pct0'),
    'IS_GENERAL': (SUP, r'^impuesto de sociedades, tipo general', 2, 'pct0'),
    'VIDA_OBRA': (SUP, r'^vida [uú]til de la furgoneta', 2, 'int'),
    'VIDA_MAQUINARIA': (SUP, r'^vida [uú]til de maquinaria', 2, 'int'),
    'AFORO': (SUP, r'^capacidad m[aá]xima por evento', 2, 'int'),
    'RAMPA_MES1': (SUP, r'^actividad del mes 1', 2, 'pct0'),
    'MESES_RAMPA': (SUP, r'^meses hasta alcanzar', 2, 'int'),
    'DIAS_COBRO': (SUP, r'^d[ií]as medios de cobro', 2, 'int'),
    'DIAS_PAGO': (SUP, r'^d[ií]as medios de pago', 2, 'int'),
    'IVA_REDUCIDO': (SUP, r'^iva repercutido de restauraci', 2, 'pct0'),
    'SMI_ANUAL': (SUP, r'^smi anual de referencia', 2, 'eur'),
    'HOLGURA_MINIMA': (SUP, r'^holgura m[ií]nima sobre el punto de equilibrio', 2, 'pct0'),
    # P&L
    'VENTAS_A1': (PYG, r'^ingresos totales', 2, 'eur'),
    'VENTAS_A2': (PYG, r'^ingresos totales', 3, 'eur'),
    'VENTAS_A3': (PYG, r'^ingresos totales', 4, 'eur'),
    'INVITADOS_EVENTO_A2': (PYG, r'^invitados por evento \(media\)', 3, 'dec1'),
    'INVITADOS_EVENTO_A3': (PYG, r'^invitados por evento \(media\)', 4, 'dec1'),
    'VENTAS_DULCES_A1': (PYG, r'^dulces y boller[ií]a', 2, 'eur'),
    'VENTAS_BEBIDAS_A1': (PYG, r'^bebidas de caf[eé] y sin caf[eé]', 2, 'eur'),
    'IVA_LINEA_DULCES': (PYG, r'^dulces y boller[ií]a', 7, 'pct0'),
    'IVA_LINEA_BEBIDAS': (PYG, r'^bebidas de caf[eé] y sin caf[eé]', 7, 'pct0'),
    'MERCANCIA_DULCES_A1': (PYG, r'^coste de mercanc[ií]a — comida', 2, 'eur'),
    'MERCANCIA_BEBIDAS_A1': (PYG, r'^coste de mercanc[ií]a — bebida', 2, 'eur'),
    'CONSUMIBLES_A1': (PYG, r'^consumibles y envases', 2, 'eur'),
    'CANON_FERIAS_A1': (PYG, r'^canon de ferias y mercados', 2, 'eur'),
    'COSTES_VARIABLES_A1': (PYG, r'^total costes variables', 2, 'eur'),
    'MARGEN_BRUTO_A1': (PYG, r'^margen bruto$', 2, 'eur'),
    'MARGEN_BRUTO_PCT_A1': (PYG, r'^margen bruto / ventas', 2, 'pct'),
    'CUOTA_BASE_ANO': (PYG, r'^cuota del almac[eé]n base$', 2, 'eur'),
    'CUOTA_BASE_PCT_A1': (PYG, r'^cuota del almac[eé]n base / ventas', 2, 'pct'),
    'PERSONAL_A1': (PYG, r'^personal \(n[oó]minas', 2, 'eur'),
    'PERSONAL_A2': (PYG, r'^personal \(n[oó]minas', 3, 'eur'),
    'PERSONAL_A3': (PYG, r'^personal \(n[oó]minas', 4, 'eur'),
    'PERSONAL_PCT_A1': (PYG, r'^coste de personal / ventas', 2, 'pct'),
    'PERSONAL_PCT_A3': (PYG, r'^coste de personal / ventas', 4, 'pct'),
    'MERCANCIA_PCT_A1': (PYG, r'^coste de mercanc[ií]a / ventas', 2, 'pct'),
    'SUMINISTROS_ANO': (PYG, r'^suministros \(', 2, 'eur'),
    'ASESORIA_ANO': (PYG, r'^asesor[ií]a', 2, 'eur'),
    'MARKETING_ANO': (PYG, r'^marketing:', 2, 'eur'),
    'MANTENIMIENTO_CAFETERA': (PYG, r'^mantenimiento de la cafetera', 2, 'eur'),
    'LICENCIAS_FERIAS': (PYG, r'^licencias y tasas de venta', 2, 'eur'),
    'AMORTIZACION_A1': (PYG, r'^amortizaci[oó]n del inmovilizado', 2, 'eur'),
    'AMORTIZACION_A2': (PYG, r'^amortizaci[oó]n del inmovilizado', 3, 'eur'),
    'AMORTIZACION_A3': (PYG, r'^amortizaci[oó]n del inmovilizado', 4, 'eur'),
    'INTERESES_A1': (PYG, r'^gastos financieros', 2, 'eur'),
    'INTERESES_A2': (PYG, r'^gastos financieros', 3, 'eur'),
    'INTERESES_A3': (PYG, r'^gastos financieros', 4, 'eur'),
    'FIJOS_A1': (PYG, r'^total costes fijos', 2, 'eur'),
    'RAI_A1': (PYG, r'^resultado antes de impuestos', 2, 'eur'),
    'RAI_A2': (PYG, r'^resultado antes de impuestos', 3, 'eur'),
    'RAI_A3': (PYG, r'^resultado antes de impuestos', 4, 'eur'),
    'IS_A1': (PYG, r'^impuesto de sociedades$', 2, 'eur'),
    'IS_A3': (PYG, r'^impuesto de sociedades$', 4, 'eur'),
    'NETO_A1': (PYG, r'^resultado neto$', 2, 'eur'),
    'NETO_A2': (PYG, r'^resultado neto$', 3, 'eur'),
    'NETO_A3': (PYG, r'^resultado neto$', 4, 'eur'),
    'NETO_PCT_A1': (PYG, r'^resultado neto / ventas', 2, 'pct'),
    'NETO_PCT_A2': (PYG, r'^resultado neto / ventas', 3, 'pct'),
    'NETO_PCT_A3': (PYG, r'^resultado neto / ventas', 4, 'pct'),
    # Inversión
    'CAPEX': (INV, r'^capex', 2, 'eur'),
    'FONDO_MANIOBRA': (INV, r'^fondo de maniobra$', 2, 'eur'),
    'INVERSION_TOTAL': (INV, r'^inversi[oó]n total', 2, 'eur'),
    'IVA_INVERSION': (INV, r'^iva soportado sobre la inversi', 2, 'eur'),
    'NECESIDAD_CAJA': (INV, r'^necesidad total de caja', 2, 'eur'),
    'FIANZA': (INV, r'^fianza del almac[eé]n base$', 2, 'eur'),
    'RENTA_PREVIA': (INV, r'^cuota del almac[eé]n base de los meses previos', 2, 'eur'),
    'IMPREVISTOS': (INV, r'^imprevistos de la compra', 2, 'eur'),
    'FURGONETA': (INV, r'^furgoneta usada', 2, 'eur'),
    'BARRA': (INV, r'^barra-carro', 2, 'eur'),
    'CAFETERA': (INV, r'^cafetera de dos grupos', 2, 'eur'),
    'MOLINOS': (INV, r'^molinos', 2, 'eur'),
    'ENERGIA': (INV, r'^energ[ií]a', 2, 'eur'),
    'AGUA': (INV, r'^agua e higiene', 2, 'eur'),
    'FRIO': (INV, r'^fr[ií]o:', 2, 'eur'),
    'MENAJE': (INV, r'^menaje', 2, 'eur'),
    'TPV': (INV, r'^tpv', 2, 'eur'),
    'MARCA_WEB': (INV, r'^marca, web', 2, 'eur'),
    'CONSTITUCION': (INV, r'^constituci[oó]n', 2, 'eur'),
    'TRAMITES': (INV, r'^declaraci[oó]n responsable', 2, 'eur'),
    'STOCK_INICIAL': (INV, r'^stock inicial', 2, 'eur'),
    'LANZAMIENTO': (INV, r'^lanzamiento', 2, 'eur'),
    'BASE_OBRA': (INV, r'^furgoneta, rotulaci[oó]n e imprevistos$', 2, 'eur'),
    'BASE_MAQUINARIA': (INV, r'^maquinaria, mobiliario y equipos', 2, 'eur'),
    # Financiación
    'PRESTAMO': (FIN, r'^pr[eé]stamo bancario$', 2, 'eur'),
    'ORIGEN_FONDOS': (FIN, r'^total origen de fondos', 2, 'eur'),
    'DIFERENCIA_FONDOS': (FIN, r'^diferencia \(origen - usos\)', 2, 'eur'),
    'DIFERENCIA_FONDOS_PCT': (FIN, r'^diferencia sobre los usos', 2, 'pct'),
    'CUOTA_ANUAL': (FIN, r'^cuota anual durante la amortizaci', 2, 'eur'),
    'DSCR_A2': (FIN, r'^dscr del a[ñn]o 2', 2, 'dec2'),
    'DSCR_A3': (FIN, r'^dscr del a[ñn]o 3', 2, 'dec2'),
    'DSCR_MIN': (FIN, r'^dscr m[ií]nimo de todo el cuadro', 2, 'dec2'),
    # Tesorería
    'SALDO_MINIMO': (TES, r'^saldo m[ií]nimo del a[ñn]o', 2, 'eur'),
    'MES_SALDO_MINIMO': (TES, r'^mes en el que la caja toca fondo', 2, 'int'),
    'SALDO_MES12': (TES, r'^saldo acumulado de caja', 13, 'eur'),
    'PAYBACK': (TES, r'^payback del proyecto', 2, 'dec1'),
    'PAYBACK_CAPEX': (TES, r'^payback sobre el capex', 2, 'dec1'),
    # Punto de equilibrio
    'COSTE_VAR_INVITADO': (PE, r'^coste variable por invitado', 2, 'eur2'),
    'MARGEN_INVITADO': (PE, r'^margen de contribuci[oó]n por invitado', 2, 'eur2'),
    'BE_INVITADOS_EVENTO': (PE, r'^invitados necesarios por evento$', 2, 'int_up'),
    'BE_INVITADOS_EVENTO_CAJA': (PE, r'^invitados necesarios por evento \(caja\)', 2, 'int_up'),
    'BE_INVITADOS_EVENTO_CAJA_A2': (PE, r'^invitados necesarios por evento \(caja\)', 3, 'int_up'),
    'BE_INVITADOS_EVENTO_CAJA_A3': (PE, r'^invitados necesarios por evento \(caja\)', 4, 'int_up'),
    'BE_INGRESOS': (PE, r'^ingresos necesarios al a[ñn]o$', 2, 'eur'),
    'BE_INGRESOS_CAJA': (PE, r'^ingresos necesarios al a[ñn]o \(caja\)', 2, 'eur'),
    'BE_INGRESOS_CAJA_A2': (PE, r'^ingresos necesarios al a[ñn]o \(caja\)', 3, 'eur'),
    'BE_INGRESOS_CAJA_A3': (PE, r'^ingresos necesarios al a[ñn]o \(caja\)', 4, 'eur'),
    'HOLGURA_CAJA': (PE, r'^holgura sobre el equilibrio de caja', 2, 'pct'),
    'HOLGURA_CAJA_A2': (PE, r'^holgura sobre el equilibrio de caja', 3, 'pct'),
    'HOLGURA_CAJA_A3': (PE, r'^holgura sobre el equilibrio de caja', 4, 'pct'),
    # Personal
    'COBERTURA_HORAS': (PER, r'^cobertura \(contratadas', 2, 'pct0'),
    'PLANTILLA_PERSONAS': (PER, r'^total plantilla', 2, 'int'),
    'JORNADAS': (PER, r'^total plantilla', 3, 'dec2'),
    'SUELDO_TITULAR': (PER, r'^propietario', 4, 'eur'),
    'BRUTO_CAMARERO': (PER, r'^camarero/a de barra', 4, 'eur'),
    'HORAS_EVENTO': (PER, r'^horas de trabajo por evento', 2, 'dec1'),
    'PERSONAS_EVENTO': (PER, r'^personas necesarias a la vez', 2, 'dec1'),
    'HORAS_NECESARIAS': (PER, r'^horas de servicio a cubrir al a[ñn]o', 2, 'int'),
    'HORAS_CONTRATADAS': (PER, r'^horas contratadas al a[ñn]o', 2, 'int'),
    'INCORPORACIONES_A2': (PER, r'^coste anual de las incorporaciones del a[ñn]o 2', 2, 'eur'),
    'INCORPORACIONES_A3': (PER, r'^coste anual de las incorporaciones del a[ñn]o 3', 2, 'eur'),
    # Escenarios (B = pesimista, D = optimista)
    'ESC_PES_INVITADOS': (ESC, r'^invitados por evento$', 2, 'int'),
    'ESC_OPT_INVITADOS': (ESC, r'^invitados por evento$', 4, 'int'),
    'ESC_PES_PRECIO': (ESC, r'^precio medio por invitado sin iva', 2, 'eur2'),
    'ESC_OPT_PRECIO': (ESC, r'^precio medio por invitado sin iva', 4, 'eur2'),
    'ESC_PES_EVENTOS': (ESC, r'^eventos al a[ñn]o', 2, 'int'),
    'ESC_OPT_EVENTOS': (ESC, r'^eventos al a[ñn]o', 4, 'int'),
    'ESC_PES_VENTAS': (ESC, r'^ingresos anuales', 2, 'eur'),
    'ESC_OPT_VENTAS': (ESC, r'^ingresos anuales', 4, 'eur'),
    'ESC_PES_NETO': (ESC, r'^resultado neto$', 2, 'eur'),
    'ESC_OPT_NETO': (ESC, r'^resultado neto$', 4, 'eur'),
}

#: Calculadora (`calculadora_caf_es.py`): token → (hoja, regex del rótulo de la columna A, columna, formato). La
#: columna es la letra fija o, entre llaves, el nombre de la cabecera de la bebida o del ejemplo.
CALC = {
    'COSTE_ESPRESSO': ('Escandallo', r'^coste de producto', '{Espresso}', 'eur2'),
    'COSTE_LATTE': ('Escandallo', r'^coste de producto', '{Latte}', 'eur2'),
    'COSTE_LATTE_ENVASE': ('Escandallo', r'^coste total con envase', '{Latte}', 'eur2'),
    'FC_LATTE': ('Escandallo', r'^food cost sin envase', '{Latte}', 'pct0'),
    'COSTE_MEDIO_BEBIDA': ('Escandallo', r'^coste medio por bebida sin envase', 'B', 'eur2'),
    'COSTE_MEDIO_ENVASE': ('Escandallo', r'^coste medio de envase', 'B', 'eur2'),
    'COSTE_MEDIO_CON_ENVASE': ('Escandallo', r'^coste medio por bebida con envase', 'B', 'eur2'),
    'CAFE_G_BEBIDA': ('Escandallo', r'^caf[eé] medio por bebida', 'B', 'dec1'),
    'FC_MEZCLA': ('Escandallo', r'^food cost de la mezcla', 'B', 'pct0'),
    'EJ_BODA_BEBIDAS': ('Ejemplos', r'^bebidas que se sirven', 'B', 'int'),
    'EJ_BODA_BARISTAS': ('Ejemplos', r'^baristas necesarios', 'B', 'int'),
    'EJ_BODA_COSTE_INVITADO': ('Ejemplos', r'^coste directo por invitado', 'B', 'eur2'),
    'EJ_BODA_PRECIO': ('Ejemplos', r'^precio por invitado sin iva', 'B', 'eur2'),
    'EJ_BODA_PRECIO_IVA': ('Ejemplos', r'^precio por invitado con iva', 'B', 'eur2'),
    'EJ_BODA_TOTAL': ('Ejemplos', r'^total del presupuesto', 'B', 'eur'),
    'EJ_BODA_MARGEN': ('Ejemplos', r'^margen bruto \(calculado', 'B', 'pct0'),
    'EJ_CONGRESO_BEBIDAS': ('Ejemplos', r'^bebidas que se sirven', 'C', 'int'),
    'EJ_CONGRESO_PRECIO': ('Ejemplos', r'^precio por invitado sin iva', 'C', 'eur2'),
    'EJ_CONGRESO_TOTAL': ('Ejemplos', r'^total del presupuesto', 'C', 'eur'),
    'EJ_FERIA_VENTA': ('Ejemplos', r'^venta prevista', 'D', 'eur'),
    'EJ_FERIA_CANON': ('Ejemplos', r'^canon del organizador \(%', 'D', 'eur'),
    'EJ_FERIA_RESULTADO': ('Ejemplos', r'^resultado de la jornada', 'D', 'eur'),
    'EJ_FERIA_TAZAS_EQUILIBRIO': ('Ejemplos', r'^tazas para cubrir', 'D', 'int'),
}


def fmt(v, f):
    return _fmt(v, f).replace('-', '−', 1) if v < 0 else _fmt(v, f)   # signo menos tipográfico


def _r(v, nd=0):
    """Redondeo comercial (mitad hacia arriba, como Excel) y no el bancario de `round()` (revisión final m2 de la
    pizzería): `round(v, 9)` quita el ruido binario del float antes de pasar a Decimal."""
    q = Decimal(1).scaleb(-nd)
    return Decimal(str(round(v, 9))).quantize(q, rounding=ROUND_HALF_UP)


def _fmt(v, f):
    if f in ('int', 'int_up'):
        n = math.ceil(v - 1e-9) if f == 'int_up' else _r(v)
        return '{:,}'.format(int(n)).replace(',', '.')
    if f == 'eur':
        return '{:,}'.format(int(_r(v))).replace(',', '.') + ' €'
    if f == 'eur2':
        return '{:,.2f}'.format(_r(v, 2)).replace(',', 'X').replace('.', ',').replace('X', '.') + ' €'
    if f in ('pct', 'pct0'):
        txt = '{:f}'.format(_r(v * 100, 1 if f == 'pct' else 0)).replace('.', ',')
        return txt + ' %'
    if f in ('dec1', 'dec2'):
        return '{:f}'.format(_r(v, 1 if f == 'dec1' else 2)).replace('.', ',')
    raise ValueError(f)


def _fila(ws, rx):
    for r in range(1, ws.max_row + 1):
        a = ws.cell(r, 1).value
        if isinstance(a, str) and re.search(rx, a.strip(), re.I):
            return r
    return None


def leer(wb):
    crudo, falta, textos = {}, [], {}
    for tok, (hoja, rx, col, f) in MAPA.items():
        ws = wb[hoja]
        r = _fila(ws, rx)
        v = ws.cell(r, col).value if r else None
        if isinstance(v, str) and tok.startswith('PAYBACK') and v.strip():   # «Más de 3 años»: texto del libro
            textos[tok] = v.strip()                                         # tal cual, para citarlo entre comillas
            textos[tok + '_TEXTO'] = v.strip()[:1].lower() + v.strip()[1:]   # en mitad de frase (finca m1)
            continue
        if not isinstance(v, (int, float)):
            falta.append('%s (%s · %s): %r' % (tok, hoja, rx, v))
            continue
        crudo[tok] = (float(v), f)
    c = dict((k, x[0]) for k, x in crudo.items())
    # derivados
    if all(k in c for k in ('CAPEX', 'FIANZA', 'RENTA_PREVIA')):        # las 14 partidas + imprevistos
        crudo['INVERSION_PARTIDAS'] = (c['CAPEX'] - c['FIANZA'] - c['RENTA_PREVIA'], 'eur')
    if 'IMPREVISTOS' in c and 'INVERSION_PARTIDAS' in crudo:            # las 14 partidas solas (SPEC D8: 42.000 €)
        crudo['INVERSION_14_PARTIDAS'] = (crudo['INVERSION_PARTIDAS'][0] - c['IMPREVISTOS'], 'eur')
    if all(k in c for k in ('VENTAS_A1', 'BE_INGRESOS')):               # holgura sobre el equilibrio CONTABLE, año 1
        crudo['HOLGURA_CONTABLE'] = (c['VENTAS_A1'] / c['BE_INGRESOS'] - 1, 'pct')
    if 'MES_SALDO_MINIMO' in c:
        textos['MES_SALDO_MINIMO_NOMBRE'] = MESES[int(round(c['MES_SALDO_MINIMO'])) - 1]
    if all(k in c for k in ('ORIGEN_FONDOS', 'NECESIDAD_CAJA')):
        crudo['EXCESO_FINANCIACION'] = (c['ORIGEN_FONDOS'] - c['NECESIDAD_CAJA'], 'eur')
    for i in (1, 2, 3):                                                   # el P&L no tiene fila de EBITDA
        ks = ('RAI_A%d' % i, 'AMORTIZACION_A%d' % i, 'INTERESES_A%d' % i)
        if all(k in c for k in ks):
            e = sum(c[k] for k in ks)
            crudo['EBITDA_A%d' % i] = (e, 'eur')
            if 'VENTAS_A%d' % i in c:
                crudo['EBITDA_PCT_A%d' % i] = (e / c['VENTAS_A%d' % i], 'pct')
    # eventos al año de equilibrio con el tamaño medio de evento del año 1 (lección M1 de la finca: dos lecturas, cada
    # una con su nombre: «X eventos del tamaño medio» o «Y invitados por evento con los eventos del caso»)
    if all(k in c for k in ('BE_INGRESOS', 'BE_INGRESOS_CAJA', 'PRECIO_INVITADO', 'INVITADOS_EVENTO')):
        evento = c['PRECIO_INVITADO'] * c['INVITADOS_EVENTO']
        crudo['VENTA_EVENTO_MEDIO'] = (evento, 'eur')
        crudo['BE_EVENTOS_ANO'] = (c['BE_INGRESOS'] / evento, 'int_up')
        crudo['BE_EVENTOS_ANO_CAJA'] = (c['BE_INGRESOS_CAJA'] / evento, 'int_up')
        for a in ('A2', 'A3'):
            if 'BE_INGRESOS_CAJA_' + a in c:
                crudo['BE_EVENTOS_CAJA_' + a] = (c['BE_INGRESOS_CAJA_' + a] / evento, 'int_up')
    # el crecimiento lo aplica el motor a los invitados por evento: en euros = eventos del mismo tamaño
    for a in ('A2', 'A3'):
        if all(k in c for k in ('VENTAS_' + a, 'VENTAS_A1', 'EVENTOS_ANO')):
            crudo['EVENTOS_EQUIVALENTES_' + a] = (c['EVENTOS_ANO'] * c['VENTAS_' + a] / c['VENTAS_A1'], 'int')
    if all(k in c for k in ('VENTAS_A1', 'CUOTA_FERIAS')):
        crudo['VENTAS_FERIAS_A1'] = (c['VENTAS_A1'] * c['CUOTA_FERIAS'], 'eur')
    # la mezcla del caso (tabla de Instrucciones y Word §6)
    total = sum(n * inv * p for _t, n, inv, p in MEZCLA)
    crudo['MEZCLA_VENTAS'] = (total, 'eur')
    crudo['MEZCLA_SERVICIOS'] = (sum(n for _t, n, _i, _p in MEZCLA), 'int')
    crudo['MEZCLA_CUOTA_FERIAS'] = (MEZCLA[-1][1] * MEZCLA[-1][2] * MEZCLA[-1][3] / total, 'pct')
    for t, n, inv, p in MEZCLA:
        crudo['TICKET_' + t.upper()] = (inv * p, 'eur')                  # factura media por evento de cada tipo
    # tabla de escalado (SPEC D13): la MISMA fuente que la calculadora y la hoja Instrucciones (`datos_escalado.py`)
    try:
        import datos_escalado as E
    except ImportError:
        falta.append('datos_escalado.py (tabla de escalado, SPEC D13)')
    else:
        crudo['TE_HORAS'] = (E.HORAS_REF, 'int')
        crudo['TE_BEBIDAS_INVITADO'] = (E.bebidas_por_invitado(E.HORAS_REF), 'dec1')
        crudo['TE_KW_ENCHUFE'] = (E.KW_ENCHUFE, 'dec1')
        for n, d in E.tabla_escalado():
            for campo, clave, f in (('BEBIDAS', 'bebidas', 'int'), ('BARISTAS', 'baristas', 'int'),
                                    ('CAFETERAS', 'estaciones', 'int'), ('MOLINOS', 'molinos', 'int'),
                                    ('KG_CAFE', 'kg_cafe', 'dec1'), ('L_LECHE', 'l_leche', 'int'),
                                    ('VASOS', 'vasos', 'int'), ('KW', 'kw', 'dec1'), ('AGUA', 'agua', 'int')):
                crudo['TE%d_%s' % (n, campo)] = (d[clave], f)
            textos['TE%d_TOMA' % n] = d['toma'][0].lower() + d['toma'][1:]
    return crudo, falta, textos


def leer_calc(path):
    """Tokens de la calculadora (opcional): caché de «Escandallo» y «Ejemplos» por rótulo."""
    wb = openpyxl.load_workbook(path, data_only=True)
    crudo, falta = {}, []
    for tok, (hoja, rx, col, f) in CALC.items():
        ws = wb[hoja]
        r = _fila(ws, rx)
        ncol = None
        if col.startswith('{'):
            nombre = col[1:-1]
            for row in ws.iter_rows(max_row=r or 1):
                for cel in row:
                    if cel.value == nombre:
                        ncol = cel.column
        else:
            ncol = openpyxl.utils.column_index_from_string(col)
        v = ws.cell(r, ncol).value if (r and ncol) else None
        if isinstance(v, (int, float)):
            crudo[tok] = (float(v), f)
        else:
            falta.append('%s (calculadora · %s · %s): %r' % (tok, hoja, rx, v))
    return crudo, falta


def prueba_iva(wb):
    """SPEC D6: las dos líneas de venta al 10 % en la columna G del P&L y el reducido de Supuestos al 10 %."""
    pyg = wb[PYG]
    out, ok = [], True
    for rx in (r'^dulces y boller[ií]a', r'^bebidas de caf[eé] y sin caf[eé]'):
        r = _fila(pyg, rx)
        g = pyg.cell(r, 7).value if r else None
        out.append('PyG!G%s = %r' % (r, g))
        ok &= isinstance(g, (int, float)) and abs(g - 0.10) < 1e-9
    red = wb[SUP].cell(_fila(wb[SUP], r'^iva repercutido de restauraci'), 2).value
    ok &= isinstance(red, (int, float)) and abs(red - 0.10) < 1e-9
    return ok, ' · '.join(out) + ' · reducido %r' % red


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--xlsx', required=True)
    ap.add_argument('--calc')
    ap.add_argument('--salida', default=SALIDA)
    a = ap.parse_args()
    wb = openpyxl.load_workbook(a.xlsx, data_only=True)
    crudo, falta, textos = leer(wb)
    falta_calc = []
    if a.calc:
        c2, falta_calc = leer_calc(a.calc)
        crudo.update(c2)
    tokens = dict((k, fmt(v, f)) for k, (v, f) in crudo.items())
    tokens.update(textos)
    json.dump({'fuente': os.path.basename(a.xlsx), 'calculadora': os.path.basename(a.calc) if a.calc else None,
               'tokens': tokens, 'valores': dict((k, v) for k, (v, _f) in crudo.items())},
              open(a.salida, 'w', encoding='utf-8'), ensure_ascii=False, indent=1, sort_keys=True)
    v = dict((k, x[0]) for k, x in crudo.items())
    iva_ok, iva_txt = prueba_iva(wb)
    gates = [
        ('DSCR mínimo ≥ 1,25', v.get('DSCR_MIN', 0) >= 1.25),
        ('saldo mínimo de tesorería > 0', v.get('SALDO_MINIMO', -1) > 0),
        ('cobertura de horas ≥ 98 %', v.get('COBERTURA_HORAS', 0) >= 0.98),
        ('holgura sobre el equilibrio de caja del año 1 ≥ 15 %', v.get('HOLGURA_CAJA', 0) >= 0.15),
        ('holgura de caja del año 2 (sin carencia) ≥ 0 % (informativa)', v.get('HOLGURA_CAJA_A2', -1) >= 0),
        ('holgura de caja del año 3 ≥ 0 % (informativa)', v.get('HOLGURA_CAJA_A3', -1) >= 0),
        ('exceso de financiación (origen - usos) ≥ 0 y < 5 %', 0 <= v.get('DIFERENCIA_FONDOS_PCT', -1) < 0.05),
        ('resultado neto a régimen (años 2 y 3) entre el 7 y el 12 %',
         all(0.07 <= v.get('NETO_PCT_A%d' % i, 1) <= 0.12 for i in (2, 3))),
        ('resultado neto del año 1 ≤ 12 %', v.get('NETO_PCT_A1', 1) <= 0.12),
        ('venta directa en ferias y mercados < 25 % (demo §2.11.9)', v.get('CUOTA_FERIAS', 1) < 0.25),
        ('venta directa = cuota de la mezcla ± 1 punto (%s frente a %s)' % (
            tokens.get('CUOTA_FERIAS'), tokens.get('MEZCLA_CUOTA_FERIAS')),
         abs(v.get('CUOTA_FERIAS', 1) - v.get('MEZCLA_CUOTA_FERIAS', 0)) <= 0.01),
        ('IVA de las dos líneas = 10 %% (%s)' % iva_txt, iva_ok),
    ]
    for tok in ('INVITADOS_EVENTO', 'PRECIO_INVITADO', 'PRECIO_INVITADO_IVA', 'EVENTOS_ANO', 'VENTAS_A1', 'VENTAS_A2',
                'VENTAS_A3', 'MARGEN_BRUTO_PCT_A1', 'MERCANCIA_PCT_A1', 'PERSONAL_A1', 'PERSONAL_PCT_A1',
                'PERSONAL_PCT_A3', 'CUOTA_BASE_PCT_A1', 'EBITDA_A1', 'EBITDA_PCT_A1', 'RAI_A1', 'NETO_A1', 'NETO_A2',
                'NETO_A3', 'NETO_PCT_A1', 'NETO_PCT_A2', 'NETO_PCT_A3', 'INVERSION_14_PARTIDAS', 'INVERSION_PARTIDAS',
                'CAPEX', 'FONDO_MANIOBRA', 'INVERSION_TOTAL', 'IVA_INVERSION', 'NECESIDAD_CAJA', 'RECURSOS_PROPIOS',
                'PRESTAMO', 'DIFERENCIA_FONDOS', 'DIFERENCIA_FONDOS_PCT', 'CUOTA_ANUAL', 'DSCR_A2', 'DSCR_A3',
                'DSCR_MIN', 'SALDO_MINIMO', 'MES_SALDO_MINIMO', 'MES_SALDO_MINIMO_NOMBRE', 'SALDO_MES12',
                'BE_INVITADOS_EVENTO', 'BE_INVITADOS_EVENTO_CAJA', 'BE_EVENTOS_ANO', 'BE_EVENTOS_ANO_CAJA',
                'BE_EVENTOS_CAJA_A2', 'BE_EVENTOS_CAJA_A3', 'HOLGURA_CONTABLE', 'HOLGURA_CAJA', 'HOLGURA_CAJA_A2',
                'HOLGURA_CAJA_A3', 'COBERTURA_HORAS', 'JORNADAS', 'PAYBACK', 'PAYBACK_CAPEX', 'CUOTA_FERIAS',
                'MEZCLA_CUOTA_FERIAS', 'ESC_PES_VENTAS', 'ESC_PES_NETO', 'ESC_OPT_VENTAS', 'ESC_OPT_NETO',
                'COSTE_ESPRESSO', 'COSTE_LATTE', 'COSTE_MEDIO_BEBIDA', 'EJ_BODA_PRECIO', 'EJ_CONGRESO_PRECIO',
                'EJ_FERIA_RESULTADO'):
        print('  %-28s %s' % (tok, tokens.get(tok, '—')))
    for nombre, ok in gates:
        print('  [%s] %s' % ('OK' if ok else 'ROJO', nombre))
    for f in falta:
        print('  FALTA', f)
    for f in falta_calc:
        print('  FALTA (calculadora, no entra en el gate)', f)
    ok_total = all(ok for _n, ok in gates) and not falta
    print('cifras_caso_es.json: %d tokens · gates del caso D9 %s' % (len(tokens), 'VERDE' if ok_total else 'ROJA'))
    sys.exit(0 if ok_total else 1)


if __name__ == '__main__':
    main()
