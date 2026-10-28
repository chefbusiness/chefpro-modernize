#!/usr/bin/env python3
"""
docx_caf_es.py — Cafetería Móvil para Eventos ES 1.0 (SPEC D15; copia de `plan-pizzeria/docx_piz_es.py`): ensambla
`plan-de-negocio-cafeteria-movil-eventos.docx` en las 131 ranuras del docx del Food Truck ES (10 Heading 1), con los
textos de `textos/caf_es_{a,b}.json` y las cifras de `cifras_caso_es.json` (tokens {ASI}, leídos de la caché del Excel,
de la calculadora y de la tabla de escalado por `cifras_caf_es.py`). Fija portada, aviso, índice, encabezados, cierre y
venta cruzada (127-129). A4; autor `AI Chef Pro`; `comments` = aichef.pro/plan-negocio-cafeteria-movil-eventos. Gates
al final: 10 Heading 1, A4, 0 tokens sin resolver, 0 prohibidos (`tandas_docx.PROHIBIDOS`, D12), «quince años» y no
«15 años», 0 no latinos, ≥ 6.000 palabras. Sin anexo (camino §4). Corre en el VPS (python-docx).

    /root/venv-guias/bin/python docx_caf_es.py --salida <dir>          # escribe <dir>/plan-de-negocio-cafeteria-movil-eventos.docx
    /root/venv-guias/bin/python docx_caf_es.py --salida <dir> --parcial # ranuras sin texto: quedan con «[PENDIENTE]»

Sesión Claude Code, 28-oct-2026.
"""
import argparse
import datetime
import json
import os
import re
import shutil
import sys

import docx
from docx.shared import Mm

AQUI = os.path.dirname(os.path.abspath(__file__))
REPO = os.path.abspath(os.path.join(AQUI, '..', '..', '..'))
FT = os.path.join(REPO, 'astro-site', 'public', 'dl', 'plan-negocio-food-truck', 'plan-de-negocio-food-truck.docx')
NOMBRE = 'plan-de-negocio-cafeteria-movil-eventos.docx'
PID = 'plan-negocio-cafeteria-movil-eventos'
sys.path.insert(0, AQUI)
sys.path.insert(0, os.path.join(os.path.dirname(AQUI), 'planes-v2_0'))
import motor                                                                       # noqa: E402
import caf_version                                                                 # noqa: E402,F401  (versión 1.0 · mes de caf_version)
import tandas_docx as T                                                            # noqa: E402

FIJOS = {
    6: 'PLAN DE NEGOCIO',
    7: 'Cafetería Móvil para Eventos',
    9: 'Plan financiero completo, análisis de mercado,\nestrategia operativa y checklist de apertura\npara montar una '
       'barra de café de especialidad para bodas, empresas y ferias',
    14: 'ChefBusiness Consultoría Gastronómica',
    16: 'Versión %s · %s' % (motor.VERSION, motor.VERSION_MES),   # la misma que estampa el motor en los xlsx
    18: 'AVISO LEGAL',
    19: 'Este documento ha sido elaborado por ChefBusiness Consultoría Gastronómica con fines informativos y '
        'orientativos. Las cifras salen del Excel del plan financiero que lo acompaña, con un caso de ejemplo que '
        'debes sustituir por el tuyo; la venta ambulante, el registro sanitario y las ordenanzas varían de una '
        'comunidad y de un ayuntamiento a otro y debes confirmarlos allí; las fuentes de los datos de mercado y de las '
        'normas se citan en el texto. Los '
        'resultados reales pueden variar. No constituye asesoramiento financiero, legal ni fiscal.',
    22: 'ÍNDICE',
    35: 'Anexos: Plan Financiero Excel + Checklist de Apertura Excel + Calculadora de Escandallo y Presupuesto Excel + '
        'Modelo de Contrato de Servicio + Guía de Equivalencias',
    124: 'Consultoría Gastronómica',
    127: '→ Kit de Tareas Recurrentes: Cafetería',
    128: '→ Kit de Escandallos Pro',
    129: '→ Pack Plantillas APPCC',
}
ENCABEZADOS = dict(zip((37, 44, 51, 59, 67, 76, 83, 91, 100, 109), (T.SECCIONES[i] for i in range(1, 11))))
INDICE = dict(zip(range(24, 34), ('%d.  %s' % (i, T.SECCIONES[i].split('. ', 1)[1]) for i in range(1, 11))))


def poner(p, texto):
    if not p.runs:
        p.add_run(texto)
        return
    p.runs[0].text = texto
    for r in p.runs[1:]:
        r.text = ''


def ensamblar(salida, parcial=False):
    cif = json.load(open(os.path.join(AQUI, 'cifras_caso_es.json'), encoding='utf-8'))['tokens']
    textos = {}
    for t in ('caf_es_a', 'caf_es_b'):
        p = os.path.join(T.TEXTOS, t + '.json')
        if os.path.exists(p):
            textos.update(json.load(open(p, encoding='utf-8')))
    destino = os.path.join(salida, NOMBRE)
    shutil.copy2(FT, destino)
    d = docx.Document(destino)
    ps = d.paragraphs
    assert len(ps) == 131, len(ps)
    pendientes = []
    for r in T.ranuras():
        i = r['i']
        t = textos.get(str(i))
        if t is None:
            if not parcial:
                raise SystemExit('falta la ranura %d (usa --parcial para un ensayo)' % i)
            pendientes.append(i)
            t = '[PENDIENTE] ' + T.SECCIONES[r['seccion']]
        t = re.sub(r'\{([A-Z0-9_]+)\}', lambda m: cif.get(m.group(1), m.group(0)), t)
        poner(ps[i], t)
    for i, t in list(FIJOS.items()) + list(ENCABEZADOS.items()) + list(INDICE.items()):
        poner(ps[i], t)
    for s in d.sections:
        s.page_width, s.page_height = Mm(210), Mm(297)
    cp = d.core_properties
    cp.author, cp.last_modified_by = 'AI Chef Pro', 'AI Chef Pro'
    cp.title = 'Plan de negocio — Cafetería Móvil para Eventos · v' + motor.VERSION
    cp.subject = 'Plan de Negocio: Cafetería Móvil para Eventos'
    cp.comments = 'aichef.pro/' + PID          # sin «generated by python-docx» (catering m13)
    cp.keywords = 'plan de negocio cafetería móvil para eventos, barra de café, coffee cart, AI Chef Pro'
    cp.created = cp.modified = datetime.datetime.now().replace(microsecond=0)
    cp.revision = 1
    d.save(destino)
    return destino, pendientes


def gates(path, parcial):
    d = docx.Document(path)
    texto = '\n'.join(p.text for p in d.paragraphs)
    err = []
    h1 = [p.text for p in d.paragraphs if p.style.name == 'Heading 1']
    if len(h1) != 10:
        err.append('Heading 1: %d (esperados 10)' % len(h1))
    s = d.sections[0]
    if abs(s.page_width - Mm(210)) > Mm(1) or abs(s.page_height - Mm(297)) > Mm(1):
        err.append('no es A4')
    sin = re.findall(r'\{[A-Z0-9_]+\}', texto)
    if sin:
        err.append('tokens sin resolver: %s' % sorted(set(sin)))
    for rx in T.PROHIBIDOS:
        m = re.search(rx, texto, 0 if rx.startswith('[a-z') else re.I)
        if m:
            err.append('prohibido /%s/: …%s…' % (rx, texto[max(0, m.start() - 40):m.end() + 40].replace('\n', ' ')))
    if re.search(r'\b15 años', texto):
        err.append('«15 años» con cifra (censo bio_vieja): escribe «quince años»')
    if not parcial and (not re.search(r'29 años', texto) or not re.search(r'quince', texto)):
        err.append('falta la bio (SPEC D15: 29 años en alta hostelería, quince en consultoría)')
    if re.search(r'[Ͱ-ϿЀ-ӿ֐-ۿ฀-๿぀-ヿ㐀-鿿가-힯]', texto):
        err.append('caracteres no latinos')
    cp = d.core_properties
    if cp.author != 'AI Chef Pro' or (cp.created and cp.created.year < 2026) or PID not in (cp.comments or ''):
        err.append('docProps: autor %r · creado %r · comments %r' % (cp.author, cp.created, cp.comments))
    palabras = len(texto.split())
    if palabras < 6000 and not parcial:
        err.append('%d palabras (< 6.000)' % palabras)
    for x in err:
        print('ERROR', x)
    print('docx: %d párrafos · %d Heading 1 · %d palabras · %s' % (
        len(d.paragraphs), len(h1), palabras, 'VERDE' if not err else 'ROJO'))
    return 1 if err else 0


if __name__ == '__main__':
    ap = argparse.ArgumentParser()
    ap.add_argument('--salida', required=True)
    ap.add_argument('--parcial', action='store_true')
    a = ap.parse_args()
    os.makedirs(a.salida, exist_ok=True)
    path, pend = ensamblar(a.salida, a.parcial)
    print('→', path, '(%d ranuras pendientes)' % len(pend) if pend else '')
    sys.exit(gates(path, a.parcial))
