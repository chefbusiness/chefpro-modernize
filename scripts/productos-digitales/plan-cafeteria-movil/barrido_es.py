#!/usr/bin/env python3
"""barrido_es.py — Cafetería Móvil para Eventos ES 1.0 (SPEC D12; copia de `plan-pizzeria/barrido_es.py`): barre los
xlsx/docx que se le pasen buscando los prohibidos de la SPEC (los 12 del research §8 —«sin licencias», comida testigo
«nunca», vaso compostable «sin plástico», «cualquier enchufe», «desde 3.000 €», facturación garantizada, servicio al 21 %,
el Ape por autovía, alcohol en la carta…— y los del brief de la pizzería —«de 22 %», «Comunitat», «15 años», «Payback» a
mitad de frase, freelance, carné de manipulador, «demanda creciente», tarifa de mercado—), los restos de los moldes (Food
Truck —la semilla—, Parrillero —el contenido clonado—, Finca) y las frases del motor sin traducir (en sala, el local,
ticket, delivery, cubierto, comensal): la lista es `tandas_docx.PROHIBIDOS` (F1-camino §3), más «Campaña» por «campana»
(`RX_CAMPANA`) y «pública» por el verbo «publica» (`RX_PUBLICA`). SIN «bollería», «leche», «generador», «furgoneta»,
«remolque» ni «aparcamiento»: aquí son el oficio. Lista celda/párrafo y texto; en las fórmulas sólo mira los literales
entre comillas. Corre en el VPS. Uso: barrido_es.py <fichero>… · exit 1 si hay coincidencias. Sesión Claude Code,
28-oct-2026."""
import os
import re
import sys

import docx
import openpyxl

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import tandas_docx as T                                                            # noqa: E402

#: «Campaña» por «campana» (motor.TILDES acentúa «campana» suelta: gotcha de la pizzería). Aquí no hay campana
#: extractora: sólo vale la campaña de apertura, lanzamiento, marketing, publicidad o Navidad.
RX_CAMPANA = r'[Cc]ampa[ñÑ]a(?! de (apertura|lanzamiento|marketing|publicidad|Navidad))'
#: «pública» por «publica» (la misma trampa de motor.TILDES con el verbo, revisión final m1 de la pizzería).
RX_PUBLICA = r'\b(nadie|se|le|les|te|nos|quien) públic[ao]n?\b|\bpúblic[ao] (precios?|tarifas?)\b'
#: Los patrones que empiezan por una clase en minúscula son sensibles a mayúsculas («Payback» a mitad de frase).
SENSIBLES = [p for p in T.PROHIBIDOS if p.startswith('[a-z')]
RX = re.compile('|'.join([p for p in T.PROHIBIDOS if p not in SENSIBLES] + [RX_CAMPANA, RX_PUBLICA]), re.I)
RX_S = re.compile('|'.join(SENSIBLES)) if SENSIBLES else None


def textos(p):
    if p.endswith('.xlsx'):
        wb = openpyxl.load_workbook(p)
        it = [(ws.title + '!' + c.coordinate, c.value) for ws in wb.worksheets for row in ws.iter_rows() for c in row
              if isinstance(c.value, str)]
        it += [('pestaña', ws.title) for ws in wb.worksheets]
        pr = wb.properties
        it += [('docProps.' + k, getattr(pr, k)) for k in ('title', 'subject', 'description', 'keywords', 'creator')
               if isinstance(getattr(pr, k), str)]
    else:
        d = docx.Document(p)
        it = [('p%d' % i, x.text) for i, x in enumerate(d.paragraphs)]
        it += [('tabla%d' % i, c.text) for i, t in enumerate(d.tables) for row in t.rows for c in row.cells]
        cp = d.core_properties
        it += [('docProps.' + k, getattr(cp, k)) for k in ('title', 'subject', 'comments', 'keywords', 'author')
               if isinstance(getattr(cp, k), str)]
    return it


def main(rutas):
    n = 0
    for p in [x for x in rutas if not x.startswith('--')]:
        for loc, t in textos(p):
            txt = ' '.join(re.findall(r'"([^"]*)"', t)) if t.startswith('=') else t
            m = RX.search(txt) or (RX_S.search(txt) if RX_S else None)
            if m:
                n += 1
                print('%s · %s · «%s» · …%s…' % (p.split('/')[-1][:28], loc, m.group(0),
                                               txt[max(0, m.start() - 50):m.end() + 50]))
    print('barrido: %d coincidencias' % n)
    return n


if __name__ == '__main__':
    sys.exit(1 if main(sys.argv[1:]) else 0)
