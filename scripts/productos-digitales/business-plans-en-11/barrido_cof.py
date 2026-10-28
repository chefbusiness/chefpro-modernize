#!/usr/bin/env python3
"""
barrido_cof.py — barrido INDEPENDIENTE de los 5 entregables del Coffee Cart Business Plan Kit (EN) en
`astro-site/public/dl/coffee-cart-business-plan/` (o en `--dir`, p. ej. el dry-run): todas las celdas de texto, los
literales de las fórmulas, los párrafos y tablas de los docx (plan y contrato) y los docProps, contra el vocabulario vetado
RX_COF_NO y las cifras prohibidas (D231: `mapas.cifras_prohibidas`), restos de español (`mapas.restos_espanol`, sin la
categoría «oficio») y caracteres no latinos; más recuentos y sha256 → `entregables_cof.json` (solo sin `--dir`). Clon del
barrido de la paella (business-plans-en-7) con `--dir`. Corre en el VPS. Sesión Claude Code, 28-oct-2026 (paso 1a).

    python3 bp11.py barrido [--dir <carpeta>]
"""
import hashlib
import json
import os
import re
import sys
from collections import OrderedDict

import docx
import openpyxl

AQUI = os.path.dirname(os.path.abspath(__file__))
REPO = os.path.abspath(os.path.join(AQUI, '..', '..', '..'))
sys.path.insert(0, os.path.join(os.path.dirname(AQUI), 'business-plans-en'))
os.environ.setdefault('BP_CONJUNTO', 'cof')
import mapas as M                                                                  # noqa: E402

DIR = os.path.join(REPO, 'astro-site', 'public', 'dl', 'coffee-cart-business-plan')
RX_LIT = re.compile(r'"((?:[^"]|"")*)"')


def piezas_xlsx(p):
    wb = openpyxl.load_workbook(p)
    out, n_f = [], 0
    for ws in wb.worksheets:
        out.append(('%s (pestaña)' % ws.title, ws.title))
        for row in ws.iter_rows():
            for c in row:
                v = c.value
                if not isinstance(v, str):
                    continue
                if c.data_type == 'f':
                    n_f += 1
                    for lit in RX_LIT.findall(v):
                        if any(ch.isalpha() for ch in lit):
                            out.append(('%s!%s (fórmula)' % (ws.title, c.coordinate), lit))
                else:
                    out.append(('%s!%s' % (ws.title, c.coordinate), v))
        for dv in ws.data_validations.dataValidation:
            for t in (dv.error, dv.errorTitle, dv.prompt, dv.promptTitle):
                if t:
                    out.append(('%s DV %s' % (ws.title, dv.sqref), t))
    pr = wb.properties
    for campo in ('title', 'subject', 'keywords', 'description', 'category', 'creator'):
        if getattr(pr, campo):
            out.append(('docProps %s' % campo, getattr(pr, campo)))
    return out, OrderedDict([('hojas', len(wb.worksheets)), ('formulas', n_f)])


def piezas_docx(p):
    d = docx.Document(p)
    out = [('párrafo %d' % i, x.text) for i, x in enumerate(d.paragraphs) if x.text.strip()]
    for t, tab in enumerate(d.tables):
        for r, row in enumerate(tab.rows):
            for c, cel in enumerate(row.cells):
                if cel.text.strip():
                    out.append(('tabla %d [%d,%d]' % (t, r, c), cel.text))
    cp = d.core_properties
    for campo in ('title', 'subject', 'keywords', 'comments', 'category', 'author'):
        if getattr(cp, campo):
            out.append(('core %s' % campo, getattr(cp, campo)))
    palabras = sum(len(x.text.split()) for x in d.paragraphs)
    return out, OrderedDict([('parrafos', len(d.paragraphs)), ('con_texto', sum(1 for x in d.paragraphs if x.text.strip())),
                             ('heading1', sum(1 for x in d.paragraphs if x.style.name == 'Heading 1')),
                             ('palabras', palabras)])


def main():
    import argparse
    ap = argparse.ArgumentParser()
    ap.add_argument('--dir', default=None, help='carpeta a barrer (por defecto, dl/coffee-cart-business-plan/)')
    a = ap.parse_args()
    carpeta = a.dir or DIR
    res, fallos, total = OrderedDict(), [], 0
    for f in sorted(os.listdir(carpeta)):
        p = os.path.join(carpeta, f)
        if f.endswith('.xlsx'):
            piezas, info = piezas_xlsx(p)
        elif f.endswith('.docx'):
            piezas, info = piezas_docx(p)
        else:
            continue
        mal = OrderedDict([('prohibido', []), ('espanol', []), ('no_latinos', [])])
        for donde, t in piezas:
            for motivo, frag in M.cifras_prohibidas(t):
                mal['prohibido'].append('%s: %s «%s»' % (donde, motivo, frag))
            r = [x for x in M.restos_espanol(t) if not x[0].startswith('oficio')]
            if r:
                mal['espanol'].append('%s: %s' % (donde, r[0][0]))
            if M.no_latinos(t):
                mal['no_latinos'].append(donde)
        total += len(piezas)
        info['piezas'] = len(piezas)
        info['sha256'] = hashlib.sha256(open(p, 'rb').read()).hexdigest()
        info.update((k, len(v)) for k, v in mal.items())
        res[f] = info
        for k, v in mal.items():
            fallos += ['%s · %s' % (f, x) for x in v]
    out = OrderedDict([('carpeta', os.path.relpath(carpeta, REPO)), ('piezas', total), ('fallos', len(fallos)),
                       ('ficheros', res)])
    if not a.dir:
        with open(os.path.join(AQUI, 'entregables_cof.json'), 'w', encoding='utf-8') as fh:
            json.dump(out, fh, ensure_ascii=False, indent=1)
            fh.write('\n')
    for x in fallos[:40]:
        print('ERROR', x)
    for f, i in res.items():
        print('%-40s %s' % (f, ' · '.join('%s %s' % (k, (v[:16] if k == 'sha256' else v)) for k, v in i.items())))
    print('barrido: %d ficheros · %d piezas · %d fallos (RX_COF_NO / cifras prohibidas / español / no latinos)'
          % (len(res), total, len(fallos)))
    return 1 if fallos else 0


if __name__ == '__main__':
    sys.exit(main())
