#!/usr/bin/env python3
"""
calculadora_cof.py — Coffee Cart Business Plan Kit (EN), SPEC D227: `coffee-cart-pricing-calculator.xlsx` «Coffee Cart Pricing
Calculator (Cost per Drink to Quote per Guest)». DUPLICADO de la calculadora ES final (`../plan-cafeteria-movil/
calculadora_caf_es.py`, `d3174a8d`, con los arreglos de su revisión: M3 sin referencias internas, m22 1,840 h, m23 rótulo del
tramo 4, política D16 de cobro y cancelación) con las mismas 4 hojas, la misma lógica y el mismo autotest, adaptada a EE. UU.:
  · «Cost per Drink» (= Escandallo): 16 insumos × 10 bebidas; café y cacao en g con precio por lb, leches y chai en fl oz con
    precio por galón, sirope en fl oz con precio por botella, envase por unidad: la columna D dice cuántas unidades de receta
    trae cada unidad de precio (453.6 g por lb, 128 fl oz por galón, 25.4 fl oz por botella de 750 ml); precio de carta SIN
    sales tax (celda para quien lo incluye); COGS con semáforo; coste de la MEZCLA (lo lee la otra hoja);
  · «Pricing Calculator» (= Calculadora): cliente Private (hosted) / Corporate (hosted) / Festival or market (guest-pay);
    política por cliente en una tabla VISIBLE (margen objetivo, cuota del organizador, retainer y saldo, cortes y tramos);
    logística por la tabla de escalado de `escalado_cof.py` (baristas, máquinas, molinos, lb de café, galones de leche,
    vasos, kW, amperios, generador en kW y agua en galones); AVISO eléctrico US (salida estándar de 110-120 V y 20 A ≈ 2.4 kW
    frente a una máquina de 220 V / 3,400 W en NEMA 6-20); precio por invitado sin sales tax (redondeo a $0.50 hacia arriba),
    con sales tax, por hora y por bebida; travel fee aparte, sin margen, con radio gratis (norma US); cobro y cancelación por
    tramos «from X to less than Y»; festival: venta (bebidas × precio o la minimum sales guarantee si es mayor), cuota del
    organizador (% + booth fee), resultado y bebidas para cubrir costes;
  · «3 Example Quotes» (boda 130 · empresa 90 · festival comunitario 150 bebidas) por las mismas fórmulas;
  · «Instructions».
US Letter, protegida sin contraseña, verdes desbloqueadas, sin símbolo de moneda en los formatos (D8), docProps D20. Corre en
el VPS (openpyxl + pycel). Sesión Claude Code, 28-oct-2026 (paso 1b).

    python3 bp11.py calculadora [--salida <dir>]            # genera + inject_cache + autotest
    python3 bp11.py calculadora [--salida <dir>] --autotest # solo el autotest
"""
import argparse
import math
import os
import re
import sys
from collections import OrderedDict

import openpyxl
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill, Protection
from openpyxl.worksheet.datavalidation import DataValidation

AQUI = os.path.dirname(os.path.abspath(__file__))
PD = os.path.dirname(AQUI)
REPO = os.path.abspath(os.path.join(AQUI, '..', '..', '..'))
sys.path.insert(0, os.path.join(PD, 'business-plans-en'))
sys.path.insert(0, AQUI)
os.environ.setdefault('BP_CONJUNTO', 'cof')
import mapas as M                                                                  # noqa: E402
import escalado_cof as DE                                                          # noqa: E402  (fuente única EN)

SLUG = 'coffee-cart-business-plan'
FICHERO = 'coffee-cart-pricing-calculator.xlsx'
TITULO = 'Coffee Cart Pricing Calculator (Cost per Drink to Quote per Guest)'
DIR_DEF = os.path.join(REPO, 'astro-site', 'public', 'dl', SLUG)
HE, H1, H2, H3 = 'Cost per Drink', 'Pricing Calculator', '3 Example Quotes', 'Instructions'
VERDE = PatternFill('solid', fgColor='E8F5E9')
GRIS = PatternFill('solid', fgColor='F2F2F2')
ROJO = PatternFill('solid', bgColor='F8CBAD', fgColor='F8CBAD')
AMBAR = PatternFill('solid', bgColor='FFE699', fgColor='FFE699')
OK = PatternFill('solid', bgColor='C6E0B4', fgColor='C6E0B4')
FUENTE = 'Calibri'
CANALES = ['Private event (hosted)', 'Corporate (hosted)', 'Festival or market (guest-pay)']
USD, ENT, PCT, DEC, USD3 = '#,##0.00', '#,##0', '0%', '#,##0.0', '#,##0.000'
FC_OK, FC_VIG, FC_ALTO = 'On target', 'Watch the cost', 'Above target'
MZ_OK, MZ_NO = 'OK: the mix adds up to 100%', 'Check the mix: it does not add up to 100%'
SERVICIOS = ['Single-use compostable cups', 'Reusable cups (venue china)']
ELEC = ['Standard 110V outlet', 'Dedicated 220V circuit', 'No outlet (generator)']
TXT_TOMA_NO = DE.TOMA_NO
TXT_AV_GEN = 'No outlet: bring the generator (minimum kW in its row) and ask the venue and the fire marshal where it can go'
TXT_AV_MAL = ('WARNING: a standard 110V outlet cannot run your machines: get a dedicated 220V circuit in writing or bring '
              'the generator')
TXT_AV_OK = 'Power covered: confirm the circuit, its voltage and amps, and the distance in writing in your contract'

# ---------------------------------------------------------------------------------------------------------------------
# Cost per Drink: insumos (nombre, precio, merma, unidades de receta por unidad de precio, origen). Sin sales tax.
# ---------------------------------------------------------------------------------------------------------------------
_LB, _GAL = DE.C['g_lb'], DE.C['oz_gal']
INSUMOS = [
    ('Coffee beans (g · $ per lb)', 14.00, 0.10, _LB,
     'Kit estimate: $14 per lb, between the national average for ground coffee ($9.33 per lb, BLS, September 2026) and '
     'specialty wholesale (about $18 per lb); 10% waste for purging and dialing in.'),
    ('Whole milk (fl oz · $ per gallon)', 4.11, 0.10, _GAL, 'National average, September 2026 (BLS: $4.106 per gallon).'),
    ('Lactose-free milk (fl oz · $ per gallon)', 7.50, 0.10, _GAL, 'Kit estimate.'),
    ('Oat milk, barista (fl oz · $ per gallon)', 13.00, 0.10, _GAL,
     'Kit estimate: about $0.80 of oat milk in an 8-oz drink (grindkit.app gives $0.60-1.10).'),
    ('Almond milk, barista (fl oz · $ per gallon)', 12.00, 0.10, _GAL, 'Kit estimate.'),
    ('Cocoa powder (g · $ per lb)', 8.00, 0.05, _LB, 'Kit estimate.'),
    ('Chai concentrate (fl oz · $ per gallon)', 20.00, 0.05, _GAL, 'Kit estimate.'),
    ('Vanilla or caramel syrup (fl oz · $ per bottle)', 9.00, 0.05, 25.4,
     'Kit estimate: a standard bottle holds 25.4 fl oz (750 ml); one pump (0.25 fl oz) costs about $0.05-0.15 '
     '(grindkit.app).'),
    ('Tea (g · $ per lb)', 30.00, 0.05, _LB, 'Free slot for your tea menu: empty in the example.'),
    ('Sugar (g · $ per lb)', 0.90, 0.0, _LB, 'Kit estimate.'),
    ('8-oz compostable hot cup (each)', 0.10, 0.03, 1.0,
     'Kit estimate: cup, lid, sleeve, napkin and stirrer come to about $0.20-0.45 per drink (grindkit.app).'),
    ('12-oz compostable hot cup (each)', 0.13, 0.03, 1.0, 'Same source.'),
    ('16-oz compostable cup (each)', 0.16, 0.03, 1.0, 'Same source.'),
    ('Lid (each)', 0.06, 0.03, 1.0, 'Same source.'),
    ('Sleeve (each)', 0.05, 0.03, 1.0, 'Same source.'),
    ('Napkin and wooden stirrer (each)', 0.03, 0.03, 1.0,
     'Same source. Some cities require compostable single-use items (Seattle, as an example): check yours.'),
]
NI = len(INSUMOS)
N_PROD = 10                                                                        # los 10 primeros son producto
IX = dict((n[0].split(' (')[0], i) for i, n in enumerate(INSUMOS))
RECETA = {
    'Espresso': dict(env=('8-oz compostable hot cup', 'Napkin and wooden stirrer')),
    'Americano': dict(env=('12-oz compostable hot cup', 'Lid', 'Sleeve', 'Napkin and wooden stirrer')),
    'Cortado': dict(env=('8-oz compostable hot cup', 'Napkin and wooden stirrer')),
    'Latte': dict(otros={'Vanilla or caramel syrup': 0.25},
                  env=('12-oz compostable hot cup', 'Lid', 'Sleeve', 'Napkin and wooden stirrer')),
    'Cappuccino': dict(otros={'Cocoa powder': 2}, env=('8-oz compostable hot cup', 'Lid', 'Sleeve', 'Napkin and wooden stirrer')),
    'Flat white': dict(env=('8-oz compostable hot cup', 'Lid', 'Sleeve', 'Napkin and wooden stirrer')),
    'Mocha': dict(otros={'Cocoa powder': 15, 'Sugar': 8},
                  env=('12-oz compostable hot cup', 'Lid', 'Sleeve', 'Napkin and wooden stirrer')),
    'Chai latte': dict(otros={'Chai concentrate': 6},
                       env=('12-oz compostable hot cup', 'Lid', 'Sleeve', 'Napkin and wooden stirrer')),
    'Hot chocolate': dict(otros={'Cocoa powder': 25, 'Sugar': 10},
                          env=('12-oz compostable hot cup', 'Lid', 'Sleeve', 'Napkin and wooden stirrer')),
    'Cold brew': dict(env=('16-oz compostable cup', 'Lid', 'Napkin and wooden stirrer')),
}
BEB = OrderedDict((b[0], b) for b in DE.BEBIDAS)
NB = len(BEB)
BCOL = [openpyxl.utils.get_column_letter(6 + j) for j in range(NB)]                # F..O
OCOL = openpyxl.utils.get_column_letter(6 + NB)                                    # P: origen del precio


def _leche_repartida(oz):
    """Reparte las fl oz de una bebida entre las cuatro leches (2 decimales); el resto del redondeo va a la entera."""
    partes = [round(oz * s, 2) for _n, s in DE.LECHES[1:]]
    return [round(oz - sum(partes), 2)] + partes


def grid():
    g = [[0.0] * NB for _ in range(NI)]
    for j, (nom, (_n, mez, dob, leche, _oz, _pvp)) in enumerate(BEB.items()):
        g[IX['Coffee beans']][j] = dob * DE.G_POR_DOBLE
        for k, v in enumerate(_leche_repartida(leche)):
            g[1 + k][j] = float(v)
        for n, q in RECETA[nom].get('otros', {}).items():
            g[IX[n]][j] = float(q)
        for n in RECETA[nom]['env']:
            g[IX[n]][j] = 1.0
    return g


ESC_IN = [('tax_menu', 'Sales tax included in your menu prices (%)', PCT, 0.0,
           'US menus usually show prices before sales tax: leave 0%. If you post round prices with the tax included (common at '
           'festivals and markets), enter your rate and the sheet takes it out.'),
          ('obj', 'Target cost of goods (% of the price before sales tax)', PCT, 0.25,
           'grindkit.app puts cost of goods at 20-30% of the revenue per drink. Use the one your financial plan targets.'),
          ('tol', 'Watch band above the target (points)', PCT, 0.05,
           'Up to this band above the target, the traffic light says "Watch the cost"; beyond it, "Above target".'),
          ('gdoble', 'Grams of coffee per double shot (g)', DEC, float(DE.G_POR_DOBLE),
           '18 g (0.63 oz) in, 36 g out (a 1:2 ratio). It moves the coffee row of every drink.')]
E = OrderedDict((k, 5 + i) for i, (k, *_r) in enumerate(ESC_IN))
E_CAB = 5 + len(ESC_IN) + 1
E_PVP, E_MIX, E_DOB = E_CAB + 1, E_CAB + 2, E_CAB + 3
E_I0 = E_DOB + 1
E_I1 = E_I0 + NI - 1
E_P1 = E_I0 + N_PROD - 1
E_R = E_I1 + 2
RES = [('leche', 'Milk and plant-based milk in total (fl oz)', DEC), ('c_prod', 'Ingredient cost', USD3),
       ('c_env', 'Cup and packaging cost', USD3), ('c_tot', 'All-in cost per drink', USD3),
       ('pvp_sin', 'Menu price before sales tax', USD), ('fc_sin', 'Cost of goods without the cup', '0.0%'),
       ('fc', 'Cost of goods with the cup', '0.0%'), ('sem', 'Cost of goods traffic light', 'General'),
       ('margen', 'Margin per drink', USD), ('pvp_obj', 'Menu price needed to hit your target', USD)]
RW = OrderedDict((k, E_R + 1 + i) for i, (k, *_r) in enumerate(RES))
E_M0 = E_R + 1 + len(RES) + 1
MZ = [('suma', 'Sum of the sales mix', PCT), ('chk', 'Mix check', 'General'),
      ('m_prod', 'Average ingredient cost per drink', USD3), ('m_env', 'Average cup and packaging cost per drink', USD3),
      ('m_tot', 'Average all-in cost per drink', USD3), ('m_cafe', 'Average coffee per drink (g)', DEC),
      ('m_leche', 'Average milk per drink (fl oz)', '0.00'), ('m_ticket', 'Average menu price per drink before sales tax', USD),
      ('m_fc', 'Cost of goods of the mix (with the cup)', '0.0%')]
MR = OrderedDict((k, E_M0 + 1 + i) for i, (k, *_r) in enumerate(MZ))
E_FIN = E_M0 + 1 + len(MZ) + 1


def esc_ref(k):
    return "'%s'!$B$%d" % (HE, MR[k])


# ---------------------------------------------------------------------------------------------------------------------
# Pricing Calculator: política por cliente, formatos, horarios
# ---------------------------------------------------------------------------------------------------------------------
FORMATOS = OrderedDict((n, v) for n, _h, v, _r in DE.BEBIDAS_POR_INVITADO)
HORARIOS = OrderedDict(DE.AJUSTE_HORARIO)
CANAL_TAB = OrderedDict([
    ('margen', ('Target gross margin on the price', PCT, (0.65, 0.60, 0.20),
                'What is left of the price after drinks, pastries, event staff, generator and the organizer fee. It pays the cart, '
                'insurance, commissary, marketing and your profit. Operators price at 2-3 times their event cost, a 50-70% '
                'gross margin (flashquotes.com); at festivals the margin is lower because you sell volume.')),
    ('canon', ('Organizer fee (% of sales)', PCT, (0.0, 0.0, 0.15),
               'Festivals and markets only (a share of sales or a flat booth fee, depending on the organizer). Kit estimate: '
               '15%. The flat booth fee goes in its own input.')),
    ('pc_senal', ('Payment: retainer at booking (% of the total)', PCT, (0.30, 0.30, 0.0),
                  'Hosted events: 30% when the client signs (US operators ask 30-50%: grindkit.app). The retainer is credited '
                  'toward the price. Festivals: no retainer.')),
    ('pc_antes', ('Payment: balance before the event (% of the total)', PCT, (0.70, 0.0, 0.0),
                  'Private events: the remaining 70%, 14 days before the event (grindkit.app).')),
    ('pc_30d', ('Payment: balance net 30 after the invoice (% of the total)', PCT, (0.0, 0.70, 0.0),
                'Corporate: the remaining 70%, net 30 from the invoice you send after the event.')),
    ('pc_acto', ('Payment: at the point of sale (% of the total)', PCT, (0.0, 0.0, 1.0),
                 'Festivals and markets: guests pay for each drink (card reader or cash).')),
    ('d1', ('Cancellation: first cutoff (days before the event)', ENT, (60, 30, 30),
            'This many days or more: tier 1. Cutoffs go down from one to the next and leave no gaps.')),
    ('d2', ('Cancellation: second cutoff (days before the event)', ENT, (30, 15, 15),
            'From here to less than the first cutoff: tier 2.')),
    ('d3', ('Cancellation: third cutoff (days before the event)', ENT, (14, 7, 7),
            'From here to less than the second cutoff: tier 3. Below this many days: tier 4. Private events: 14 days, the day '
            'the balance is due.')),
    ('cn1', ('Cancellation tier 1 (% of the total you keep)', PCT, (0.10, 0.10, 0.0),
             'Hosted events: 10% of the total. Festivals: not used.')),
    ('cn2', ('Cancellation tier 2 (% of the total you keep)', PCT, (0.30, 0.30, 0.0), 'Hosted events: 30% of the total.')),
    ('cn3', ('Cancellation tier 3 (% of the total you keep)', PCT, (0.60, 0.60, 0.0), 'Hosted events: 60% of the total.')),
    ('cn4', ('Cancellation tier 4 (% of the total you keep)', PCT, (1.0, 1.0, 0.0),
             'Late cancellation or no-show: 100% minus the costs you save (coffee, milk and cups you do not use, staff and '
             'travel). At festivals, the organizer\'s vendor terms apply.')),
])
LABELS = OrderedDict([('lab1', 'Tier 1'), ('lab2', 'Tier 2'), ('lab3', 'Tier 3'), ('lab4', 'Tier 4')])

ENTRADAS = [
    ('canal', 'Client type', '@', CANALES[0],
     'Private event, Corporate or Festival or market: they change the target margin, the organizer fee, the payments and '
     'the cancellation tiers (table at the bottom). At a festival, "guests" are the drinks you expect to sell.'),
    ('guests', 'Expected guests (at a festival or market, expected drinks)', ENT, 130, 'Whole number of guests you serve.'),
    ('formato', 'Service length (sets the drinks per guest)', '@', 'Over 2 and up to 4 hours',
     'Drinks per guest by the hours the bar is open: table at the bottom (Kit estimate from US ratios). Not used at '
     'festivals.'),
    ('horario', 'Time of day', '@', 'Evening (after dinner)',
     'Morning +25%, evening after dinner −10% (Kit estimate from US ratios). It adjusts the drinks per guest. Not used at '
     'festivals.'),
    ('serv', 'Served in', '@', SERVICIOS[0],
     'Single-use compostable cups add the cup, lid and sleeve to the cost. Reusable cups (the venue\'s china) carry no cup '
     'cost; some venues and cities ask for them.'),
    ('hours', 'Service hours (bar open)', DEC, 3, 'Used for the drinks per hour and for staffing.'),
    ('h_extra', 'Hours per person outside service (loading, travel, setup and cleanup)', DEC, 5,
     'Kit estimate: a 2-hour event takes a 7-9 hour shift (grindkit.app). Use your own.'),
    ('dulces', 'Pastries per guest (your cost, from a licensed bakery)', USD, 0.75,
     'What the pastries you add to the package cost you. 0 if you do not offer them.'),
    ('rentals', 'Rentals billed to the client (tent, tables, extra cups)', USD, 0.0,
     'What the rental company charges you for this event.'),
    ('elec', 'Power at the venue', '@', ELEC[1],
     'A standard 110-120V outlet on a 20-amp circuit gives about 2.4 kW. A 2-group machine at 220V draws about 3.4 kW '
     '(NEMA 6-20 plug, 15-17 A): it needs a dedicated 220V circuit or a generator. A compact machine can run on 110V.'),
    ('gen_cost', 'Generator: cost per event (fuel and wear, only without an outlet)', USD, 60.0,
     'Kit estimate. Only added if you choose "No outlet (generator)".'),
    ('kw_maq', 'Power draw of your 2-group machine (kW)', DEC, DE.KW_MAQUINA,
     'The example machine: 3,400 W at 220V. Check the rating plate of yours.'),
    ('wage', 'Hourly wage of a barista', USD, 20.0,
     'Kit estimate: between the event barista ($19 per hour) and the lead barista ($20) of the financial plan; '
     'Flashquotes models $25 per hour. Use your own.'),
    ('ss', 'Payroll taxes on wages (%)', PCT, 0.10,
     'FICA (7.65%) plus federal and state unemployment taxes, as in the financial plan. Workers\' comp is in your '
     'insurance.'),
    ('km', 'Miles one way (base to venue)', ENT, 25,
     'Charged round trip on a separate line, with no margin (at a festival it is your own cost).'),
    ('free_mi', 'Free travel radius (miles one way)', ENT, 20,
     'Many operators include the first 20-30 miles and charge beyond them (grindkit.app).'),
    ('per_km', 'Travel rate per mile (charged round trip)', USD, 1.25,
     'flashquotes.com shows $1.25 per mile round trip. Use your own.'),
    ('iva', 'Sales tax rate', PCT, 0.08,
     'Example rate applied to the whole invoice. Many states tax catering, including the staff and extras (some tax the travel fee too); at '
     'festivals you collect it on each drink. Look up your state and city.'),
    ('tkt', 'Festival: average price per drink you will charge (before sales tax)', USD, 6.50,
     'Kit estimate: about $6.50 per drink. See the average menu price of your mix on the "%s" tab. Only used for '
     'festivals and markets.' % HE),
    ('canon_fijo', 'Festival: flat booth fee', USD, 0.0, 'A flat fee per spot, if the organizer charges one. Festivals only.'),
    ('f_min', 'Festival: minimum sales guarantee from the organizer (0 if none)', USD, 0.0,
     'Some organizers guarantee a minimum and pay the difference if your drink sales fall short. Festivals only.'),
    ('dias', 'Days before the event if the client cancels', ENT, 45,
     'Pick the days and the sheet shows the tier and what you keep.'),
]
K = OrderedDict((k, 5 + i) for i, (k, *_r) in enumerate(ENTRADAS))
F_CALC = 5 + len(ENTRADAS) + 1

SALIDAS = [
    ('h_esc', 'LOGISTICS (scaling table)', None, None, ''),
    ('beb_inv', 'Drinks per guest for this service length', '0.00', '{t_beb_inv}', 'Looked up in the table at the bottom.'),
    ('ajuste', 'Time-of-day adjustment', PCT, '{t_ajuste}', 'Looked up in the table at the bottom.'),
    ('bebidas', 'Drinks served', ENT, 'IF({is_feria},{guests},ROUNDUP(ROUND({guests}*{beb_inv}*(1+{ajuste}),2),0))',
     'Guests × drinks per guest × (1 + adjustment), rounded up. At a festival, the expected drinks.'),
    ('por_hora', 'Drinks per hour (average)', DEC, '{bebidas}/{hours}', ''),
    ('baristas', 'Baristas needed', ENT, 'MAX(1,ROUNDUP(ROUND({por_hora}*{r_pico}/{r_cap_barista},2),0))',
     'Drinks per hour × peak factor ÷ what one barista makes, rounded up.'),
    ('estaciones', '2-group machines (stations)', ENT, 'ROUNDUP({baristas}/{r_bar_maq},0)',
     'Two baristas per machine; a second station comes in when the peak does not fit on one.'),
    ('grupos', 'Groups in total', ENT, '{estaciones}*{r_grupos_maq}', ''),
    ('molinos', 'Grinders', ENT, '{estaciones}+{r_molinos_extra}', 'One per station plus a spare for decaf.'),
    ('lb_cafe', 'Coffee beans (lb)', DEC, 'ROUNDUP({bebidas}*{esc_cafe}/{c_g_lb}+{r_purga},1)',
     'Drinks × average coffee of your mix ("%s" tab) plus purging and dialing in.' % HE),
    ('gal_leche', 'Milk and plant-based milks (gal)', ENT, 'ROUNDUP(ROUND({bebidas}*{esc_leche}/{c_oz_gal},2),0)',
     'Drinks × average milk of your mix.'),
    ('vasos', 'Single-use cups or reusable cups', ENT,
     'IF({is_loza},ROUNDUP(ROUND({guests}*{r_loza_inv},2),0),ROUNDUP(ROUND({bebidas}*(1+{r_reserva}),2),0))',
     'Single-use: drinks plus the spares. Reusable: cups per guest with spares while they are washed.'),
    ('agua', 'Fresh water (gal)', ENT, 'ROUNDUP(ROUND({bebidas}*{r_agua_beb}+{r_agua_base},2),0)',
     'Bring a gray-water tank at least as large: you empty it at the commissary, never at the venue.'),
    ('h_ele', 'POWER (warning for the contract)', None, None, ''),
    ('kw', 'Power draw of the machines (kW)', DEC, '{estaciones}*{kw_maq}', 'Stations × kW of your machine.'),
    ('amp', 'Current on a 220V circuit (A)', DEC, '{kw}*1000/{r_v_maquina}', 'kW × 1,000 ÷ circuit voltage.'),
    ('kw_enchufe', 'Power of a standard outlet (kW)', '0.00', '{r_amp_enchufe}*{r_v_enchufe}/1000',
     'A 20-amp circuit at 120V gives about 2.4 kW. Grinders, fridge and lights need their own margin.'),
    ('gen_kw', 'Minimum inverter generator (kW running)', DEC, 'ROUNDUP(ROUND({kw}*{r_factor_kw},2),1)',
     'kW × sizing factor, rounded up: one 3.4 kW machine needs a 5.5 kW (5,500 W) inverter generator.'),
    ('toma', 'Power needed', 'General', 'IF({kw}<={kw_enchufe},$A${e1},"%s")' % TXT_TOMA_NO,
     'A compact machine can run on a standard outlet; a 220V 2-group machine cannot.'),
    ('aviso', 'POWER WARNING', 'General',
     'IF({elec}=$A${e3},"%s",IF(AND({kw}>{kw_enchufe},{elec}=$A${e1}),"%s","%s"))' % (TXT_AV_GEN, TXT_AV_MAL, TXT_AV_OK),
     'The kit\'s contract asks for the circuit in writing: voltage, amps and distance.'),
    ('h_cos', 'DIRECT COST', None, None, ''),
    ('c_beb', 'Drink cost', USD, '{bebidas}*IF({is_loza},{esc_sin},{esc_con})',
     'Drinks × average cost of the mix from the "%s" tab (with the cup if single-use).' % HE),
    ('c_dulces', 'Pastry cost', USD, '{dulces}*{guests}', ''),
    ('c_staff', 'Staff cost (wages + payroll taxes)', USD, '{baristas}*({hours}+{h_extra})*{wage}*(1+{ss})',
     'Baristas × (service hours + hours outside service) × hourly cost. If you work the bar yourself, count yourself: it '
     'is your opportunity cost.'),
    ('c_gen', 'Generator', USD, 'IF({elec}=$A${e3},{gen_cost},0)', ''),
    ('c_rent', 'Rentals billed to the client', USD, '{rentals}', ''),
    ('c_desp_f', 'Travel to the festival (your own cost)', USD, 'IF({is_feria},{km}*2*{per_km},0)',
     'Nobody pays your travel at a festival: it is part of the cost. For hosted events it goes on its own line.'),
    ('c_total', 'Total direct cost', USD, '{c_beb}+{c_dulces}+{c_staff}+{c_gen}+{c_rent}+{c_desp_f}', ''),
    ('c_guest', 'Direct cost per guest', USD, '{c_total}/{guests}', ''),
    ('h_pre', 'PRICE AND QUOTE', None, None, ''),
    ('margen', 'Target margin for this client type', PCT, '{t_margen}', 'From the client table at the bottom.'),
    ('canon', 'Organizer fee for this client type', PCT, '{t_canon}', 'From the client table at the bottom. Festivals only.'),
    ('price', 'PRICE PER GUEST BEFORE SALES TAX (MINIMUM)', USD, 'ROUNDUP({c_total}/(1-{margen}-{canon})/{guests}*2,0)/2',
     'Cost ÷ (1 − margin − fee) ÷ guests, rounded up to the next $0.50: the lowest price that reaches your target margin. '
     'At a festival it is the lowest price per drink that covers cost, fee and margin: compare it with yours.'),
    ('price_iva', 'Price per guest incl. sales tax', USD, '{price}*(1+{iva})',
     'For reference: US quotes show the price before tax and add the sales tax on the invoice.'),
    ('subtotal', 'Coffee service subtotal before sales tax', USD, '{price}*{guests}', 'Price per guest × guests.'),
    ('price_h', 'Same price per service hour (before sales tax)', USD, '{subtotal}/{hours}',
     'Subtotal ÷ service hours: to compare with an hourly rate.'),
    ('price_taza', 'Same price per drink (before sales tax)', USD, '{subtotal}/{bebidas}',
     'Subtotal ÷ drinks: to compare with a per-drink price.'),
    ('desp', 'Travel fee (round trip beyond the free radius, no margin)', USD,
     'IF({is_feria},0,MAX(0,{km}-{free_mi})*2*{per_km})', 'Its own line on the quote.'),
    ('base', 'Taxable amount', USD, '{subtotal}+{desp}', ''),
    ('iva_amt', 'Sales tax', USD, '{base}*{iva}', ''),
    ('total', 'QUOTE TOTAL (incl. sales tax)', USD, '{base}+{iva_amt}', ''),
    ('h_cob', 'PAYMENTS', None, None, ''),
    ('cb1', 'Payment 1: retainer at booking', USD, '{total}*{t_pc_senal}',
     'Credited toward the price. Send a receipt for every payment, the retainer too.'),
    ('cb2', 'Payment 2: balance before the event', USD, '{total}*{t_pc_antes}', 'Private events: 14 days before.'),
    ('cb3', 'Payment 3: balance net 30 after the invoice', USD, '{total}*{t_pc_30d}', 'Corporate.'),
    ('cb4', 'Payment 4: at the point of sale', USD, '{total}*{t_pc_acto}', 'Festivals and markets.'),
    ('canon_amt', 'Organizer fee on the subtotal', USD, '{subtotal}*{canon}', ''),
    ('profit', 'Gross margin in dollars', USD, '{subtotal}-{c_total}-{canon_amt}', 'On the subtotal, without the travel fee.'),
    ('margin_calc', 'Gross margin (calculated)', '0.0%', 'IF({subtotal}=0,0,{profit}/{subtotal})',
     'Comes from the costs, never typed. It stays at or above the target because the price is rounded up.'),
    ('h_can', 'CANCELLATION (tiers with no gaps)', None, None, ''),
    ('tramo', 'Tier for the days you chose', 'General',
     'IF({dias}>={t_d1},{t_lab1},IF({dias}>={t_d2},{t_lab2},IF({dias}>={t_d3},{t_lab3},{t_lab4})))',
     'Looks at the "Days before the event" input.'),
    ('ret_pct', 'Share you keep in that tier', PCT,
     'IF({dias}>={t_d1},{t_cn1},IF({dias}>={t_d2},{t_cn2},IF({dias}>={t_d3},{t_cn3},{t_cn4})))', ''),
    ('ret', 'Amount you keep if the client cancels that day (in the last tier, before you deduct the costs you save)', USD,
     '{total}*{ret_pct}', ''),
    ('can1', 'Amount kept: tier 1 (the earliest)', USD, '{total}*{t_cn1}', 'The cutoffs of each client type are in the table at '
                                                                         'the bottom.'),
    ('can2', 'Amount kept: tier 2', USD, '{total}*{t_cn2}', ''),
    ('can3', 'Amount kept: tier 3', USD, '{total}*{t_cn3}', ''),
    ('can4', 'Amount kept: tier 4 (the latest, or a no-show)', USD, '{total}*{t_cn4}', ''),
    ('h_fer', 'FESTIVAL OR MARKET: THE DAY AHEAD (only if the client type is a festival)', None, None, ''),
    ('c_var', 'Variable cost per drink (drink + pastry)', USD, '({c_beb}+{c_dulces})/{guests}', ''),
    ('f_venta', 'Expected sales (drinks × average price, or the minimum sales guarantee if higher; before sales tax)', USD,
     'IF({is_feria},MAX({guests}*{tkt},{f_min}),0)', ''),
    ('f_canon', 'Organizer fee (% of sales + flat booth fee)', USD, 'IF({is_feria},{f_venta}*{canon}+{canon_fijo},0)', ''),
    ('f_result', 'Result of the day (sales − fee − cost)', USD, 'IF({is_feria},{f_venta}-{f_canon}-{c_total},0)',
     'Before income tax and the fixed costs of the business (cart, insurance, commissary).'),
    ('f_eq', 'Drinks needed to cover the costs of the day', ENT,
     'IF({is_feria},IF({tkt}*(1-{canon})-{c_var}<=0,0,ROUNDUP(({c_staff}+{c_gen}+{c_rent}+{c_desp_f}+{canon_fijo})/'
     '({tkt}*(1-{canon})-{c_var}),0)),0)',
     'Costs of the day that do not depend on the drinks ÷ what each drink leaves after the fee (without the guarantee).'),
]
S = OrderedDict()
_fila = F_CALC
for _s in SALIDAS:
    _fila += 1
    if _s[3] is not None:
        S[_s[0]] = _fila
F_FIN_SAL = _fila

F_REGLAS = F_FIN_SAL + 2
RG = OrderedDict((k, F_REGLAS + 1 + i) for i, (k, *_r) in enumerate(DE.REGLAS))
F_CONV = F_REGLAS + 1 + len(DE.REGLAS)
CV = OrderedDict((k, F_CONV + 1 + i) for i, (k, *_r) in enumerate(DE.CONVERSIONES))
F_FORM = F_CONV + 1 + len(DE.CONVERSIONES) + 1
F_HOR = F_FORM + 1 + len(FORMATOS) + 1
F_SERV = F_HOR + 1 + len(HORARIOS) + 1
F_ELEC = F_SERV + 1 + len(SERVICIOS) + 1
F_CAN = F_ELEC + 1 + len(ELEC) + 1
CT = OrderedDict((k, F_CAN + 1 + i) for i, k in enumerate(list(CANAL_TAB) + list(LABELS)))
F_ESC = F_CAN + 1 + len(CT) + 1
ESC_HORAS = F_ESC + 1
ESC_F0 = ESC_HORAS + 1
F_ULT = ESC_F0 + len(DE.FILAS_TABLA)

_BASE = OrderedDict((k, v) for k, _r, _f2, v, _n in ENTRADAS)
EJEMPLOS = OrderedDict([
    ('Wedding coffee bar after dinner', dict(_BASE)),
    ('Corporate morning meeting', dict(
        _BASE, canal=CANALES[1], guests=90, formato='Up to 2 hours', horario='Morning (breakfast and mid-morning coffee)',
        hours=2, h_extra=5, dulces=0.0, km=15, dias=20)),
    ('Community festival, guest-pay', dict(
        _BASE, canal=CANALES[2], guests=150, formato='Up to 2 hours', horario='Morning (breakfast and mid-morning coffee)',
        hours=4, h_extra=3, dulces=0.0, elec=ELEC[2], km=15, f_min=800.0, dias=45)),
])
EJ_NOTAS = OrderedDict([
    ('Wedding coffee bar after dinner', 'Private event, 130 guests (the wedding of the financial plan\'s example mix), a 3-hour '
     'bar after dinner in compostable cups, with a $0.75 pastry per guest and a dedicated 220V circuit. It quotes the '
     'lowest price that reaches the 65% target margin; the 30% retainer at booking and the 70% balance 14 days before. With '
     '45 days to go, a cancellation falls in tier 2 and you keep 30%.'),
    ('Corporate morning meeting', 'Corporate, 90 guests (the company event of the example mix), a 2-hour bar in the morning '
     '(+25% drinks), compostable cups and a dedicated circuit. The quote is before sales tax; 30% retainer at signing and '
     'the 70% balance net 30 after the invoice. With 20 days to go, the tier is the second and you keep 30%.'),
    ('Community festival, guest-pay', 'Festival or market (a community festival): 150 drinks over 4 hours in the morning, compostable cups, a generator '
     'and a 15% organizer fee, with an $800 minimum sales guarantee. The "price per guest" is the lowest price per drink; at '
     '$6.50 a drink the day leaves about ${res} before fixed costs, and, without the guarantee, below {eq} drinks you would '
     'lose money: the $800 guarantee covers that risk.'),
])
OBJETIVO_EJ = OrderedDict([('Wedding coffee bar after dinner', 15.0), ('Corporate morning meeting', 12.0)])


# ---------------------------------------------------------------------------------------------------------------------
# Cálculo en Python (espejo de las fórmulas) — la fuente del autotest
# ---------------------------------------------------------------------------------------------------------------------
_ceil = DE._ceil


def calcular_escandallo(g=None, ing=None, mezcla=None, pvp=None, tax_menu=0.0, obj=0.25, tol=0.05):
    g = g or grid()
    ing = ing or [(p, m, u) for _n, p, m, u, _o in INSUMOS]
    mezcla = mezcla or [b[1] / 100.0 for b in BEB.values()]
    pvp = pvp or [b[5] for b in BEB.values()]
    neto = [p / (1 - m) / u for p, m, u in ing]
    r = OrderedDict()
    r['neto'] = neto
    r['bebidas'] = []
    for j in range(NB):
        d = OrderedDict()
        d['leche'] = sum(g[i][j] for i in range(1, 5))
        d['c_prod'] = sum(g[i][j] * neto[i] for i in range(N_PROD))
        d['c_env'] = sum(g[i][j] * neto[i] for i in range(N_PROD, NI))
        d['c_tot'] = d['c_prod'] + d['c_env']
        d['pvp_sin'] = pvp[j] / (1 + tax_menu)
        d['fc_sin'] = d['c_prod'] / d['pvp_sin']
        d['fc'] = d['c_tot'] / d['pvp_sin']
        d['sem'] = FC_OK if d['fc'] <= obj else (FC_VIG if d['fc'] <= obj + tol else FC_ALTO)
        d['margen'] = d['pvp_sin'] - d['c_tot']
        d['pvp_obj'] = d['c_tot'] / obj * (1 + tax_menu)
        r['bebidas'].append(d)
    suma = sum(mezcla)
    wp = lambda f: 0 if suma == 0 else sum(mezcla[j] * f(j) for j in range(NB)) / suma      # noqa: E731
    r['mz'] = OrderedDict(
        suma=suma, chk=MZ_OK if round(suma, 2) == 1 else MZ_NO,
        m_prod=wp(lambda j: r['bebidas'][j]['c_prod']), m_env=wp(lambda j: r['bebidas'][j]['c_env']),
        m_cafe=wp(lambda j: g[0][j]), m_leche=wp(lambda j: r['bebidas'][j]['leche']),
        m_ticket=wp(lambda j: r['bebidas'][j]['pvp_sin']))
    r['mz']['m_tot'] = r['mz']['m_prod'] + r['mz']['m_env']
    r['mz']['m_fc'] = r['mz']['m_tot'] / r['mz']['m_ticket']
    return r


ESC = calcular_escandallo()


def calcular(e, esc=None):
    esc = esc or ESC
    mz = esc['mz']
    r = OrderedDict()
    ic = CANALES.index(e['canal'])
    t = lambda k: CANAL_TAB[k][2][ic]                                             # noqa: E731
    reg = dict(DE.R)
    feria = e['canal'] == CANALES[2]
    loza = e['serv'] == SERVICIOS[1]
    r['beb_inv'] = FORMATOS[e['formato']]
    r['ajuste'] = HORARIOS[e['horario']]
    r['bebidas'] = e['guests'] if feria else DE._cu(e['guests'] * r['beb_inv'] * (1 + r['ajuste']))
    r['por_hora'] = r['bebidas'] / float(e['hours'])
    r['baristas'] = max(1, DE._cu(r['por_hora'] * reg['pico'] / reg['cap_barista']))
    r['estaciones'] = int(_ceil(r['baristas'] / float(reg['bar_maq'])))
    r['grupos'] = r['estaciones'] * reg['grupos_maq']
    r['molinos'] = r['estaciones'] + reg['molinos_extra']
    r['lb_cafe'] = _ceil(r['bebidas'] * mz['m_cafe'] / DE.C['g_lb'] + reg['purga'], 1)
    r['gal_leche'] = DE._cu(r['bebidas'] * mz['m_leche'] / DE.C['oz_gal'])
    r['vasos'] = DE._cu(e['guests'] * reg['loza_inv']) if loza else DE._cu(r['bebidas'] * (1 + reg['reserva']))
    r['agua'] = DE._cu(r['bebidas'] * reg['agua_beb'] + reg['agua_base'])
    r['kw'] = r['estaciones'] * e['kw_maq']
    r['amp'] = r['kw'] * 1000.0 / reg['v_maquina']
    r['kw_enchufe'] = reg['amp_enchufe'] * reg['v_enchufe'] / 1000.0
    r['gen_kw'] = DE._cu(r['kw'] * reg['factor_kw'], 1)
    r['toma'] = ELEC[0] if r['kw'] <= r['kw_enchufe'] else TXT_TOMA_NO
    r['aviso'] = (TXT_AV_GEN if e['elec'] == ELEC[2] else
                  (TXT_AV_MAL if (r['kw'] > r['kw_enchufe'] and e['elec'] == ELEC[0]) else TXT_AV_OK))
    r['c_beb'] = r['bebidas'] * (mz['m_prod'] if loza else mz['m_tot'])
    r['c_dulces'] = e['dulces'] * e['guests']
    r['c_staff'] = r['baristas'] * (e['hours'] + e['h_extra']) * e['wage'] * (1 + e['ss'])
    r['c_gen'] = e['gen_cost'] if e['elec'] == ELEC[2] else 0
    r['c_rent'] = e['rentals']
    r['c_desp_f'] = e['km'] * 2 * e['per_km'] if feria else 0
    r['c_total'] = r['c_beb'] + r['c_dulces'] + r['c_staff'] + r['c_gen'] + r['c_rent'] + r['c_desp_f']
    r['c_guest'] = r['c_total'] / e['guests']
    r['margen'] = t('margen')
    r['canon'] = t('canon')
    r['price'] = _ceil(r['c_total'] / (1 - r['margen'] - r['canon']) / e['guests'] * 2) / 2
    r['price_iva'] = r['price'] * (1 + e['iva'])
    r['subtotal'] = r['price'] * e['guests']
    r['price_h'] = r['subtotal'] / e['hours']
    r['price_taza'] = r['subtotal'] / r['bebidas']
    r['desp'] = 0 if feria else max(0, e['km'] - e['free_mi']) * 2 * e['per_km']
    r['base'] = r['subtotal'] + r['desp']
    r['iva_amt'] = r['base'] * e['iva']
    r['total'] = r['base'] + r['iva_amt']
    r['cb1'] = r['total'] * t('pc_senal')
    r['cb2'] = r['total'] * t('pc_antes')
    r['cb3'] = r['total'] * t('pc_30d')
    r['cb4'] = r['total'] * t('pc_acto')
    r['canon_amt'] = r['subtotal'] * r['canon']
    r['profit'] = r['subtotal'] - r['c_total'] - r['canon_amt']
    r['margin_calc'] = 0 if r['subtotal'] == 0 else r['profit'] / r['subtotal']
    d1, d2, d3 = t('d1'), t('d2'), t('d3')
    labs = etiquetas(d1, d2, d3)
    k = 0 if e['dias'] >= d1 else (1 if e['dias'] >= d2 else (2 if e['dias'] >= d3 else 3))
    r['tramo'] = labs[k]
    r['ret_pct'] = t('cn%d' % (k + 1))
    r['ret'] = r['total'] * r['ret_pct']
    for i in (1, 2, 3, 4):
        r['can%d' % i] = r['total'] * t('cn%d' % i)
    r['c_var'] = (r['c_beb'] + r['c_dulces']) / e['guests']
    r['f_venta'] = max(e['guests'] * e['tkt'], e['f_min']) if feria else 0
    r['f_canon'] = r['f_venta'] * r['canon'] + e['canon_fijo'] if feria else 0
    r['f_result'] = r['f_venta'] - r['f_canon'] - r['c_total'] if feria else 0
    den = e['tkt'] * (1 - r['canon']) - r['c_var']
    r['f_eq'] = 0 if (not feria or den <= 0) else int(_ceil((r['c_staff'] + r['c_gen'] + r['c_rent'] + r['c_desp_f']
                                                             + e['canon_fijo']) / den))
    return r


def etiquetas(d1, d2, d3):
    f = lambda x: ('%g' % x)                                                      # noqa: E731
    return ['%s days or more' % f(d1), 'From %s to less than %s days' % (f(d2), f(d1)),
            'From %s to less than %s days' % (f(d3), f(d2)), 'Less than %s days' % f(d3)]


# ---------------------------------------------------------------------------------------------------------------------
# Fórmulas
# ---------------------------------------------------------------------------------------------------------------------
def _rng(col, f1, f2, hoja=H1):
    return "'%s'!$%s$%d:$%s$%d" % (hoja, col, f1, col, f2)


def refs(col):
    ref = dict((k, '$%s$%d' % (col, f)) for k, f in K.items())
    ref.update((k, '$%s$%d' % (col, f)) for k, f in S.items())
    ref['t_beb_inv'] = 'INDEX(%s,MATCH(%s,%s,0))' % (_rng('B', F_FORM + 1, F_FORM + len(FORMATOS)), ref['formato'],
                                                     _rng('A', F_FORM + 1, F_FORM + len(FORMATOS)))
    ref['t_ajuste'] = 'INDEX(%s,MATCH(%s,%s,0))' % (_rng('B', F_HOR + 1, F_HOR + len(HORARIOS)), ref['horario'],
                                                    _rng('A', F_HOR + 1, F_HOR + len(HORARIOS)))
    for k, f in RG.items():
        ref['r_' + k] = "'%s'!$B$%d" % (H1, f)
    for k, f in CV.items():
        ref['c_' + k] = "'%s'!$B$%d" % (H1, f)
    cab = "'%s'!$B$%d:$D$%d" % (H1, F_CAN, F_CAN)
    for k, f in CT.items():
        ref['t_' + k] = "INDEX('%s'!$B$%d:$D$%d,MATCH(%s,%s,0))" % (H1, f, f, ref['canal'], cab)
    ref['is_feria'] = "%s='%s'!$D$%d" % (ref['canal'], H1, F_CAN)
    ref['is_loza'] = "%s='%s'!$A$%d" % (ref['serv'], H1, F_SERV + 2)
    ref['e1'] = str(F_ELEC + 1)
    ref['e3'] = str(F_ELEC + 3)
    for k, c in (('esc_sin', 'm_prod'), ('esc_con', 'm_tot'), ('esc_cafe', 'm_cafe'), ('esc_leche', 'm_leche')):
        ref[k] = esc_ref(c)
    ref['kw_enchufe'] = '$%s$%d' % (col, S['kw_enchufe'])
    return ref


def a_excel(expr, ref):
    e = re.sub(r'\$A\$\{(e\d)\}', lambda m: "'%s'!$A$%s" % (H1, ref[m.group(1)]), expr)
    return re.sub(r'\{(\w+)\}', lambda m: ref[m.group(1)], e)


def _f(ws, coord, v, b=False, sz=11, i=False, fmt=None, wrap=False, fill=None, locked=True):
    c = ws[coord]
    c.value = v
    c.font = Font(name=FUENTE, size=sz, bold=b, italic=i)
    if fmt:
        c.number_format = fmt
    c.alignment = Alignment(wrap_text=wrap, vertical='top')
    if fill:
        c.fill = fill
    c.protection = Protection(locked=locked)
    return c


def _nota(ws, fila, texto, desde='C', hasta='F', chars=105.0):
    if not texto:
        return
    _f(ws, '%s%d' % (desde, fila), texto, wrap=True)
    ws.merge_cells('%s%d:%s%d' % (desde, fila, hasta, fila))
    ws.row_dimensions[fila].height = max(16, 15 * math.ceil(len(texto) / chars))


def _imprimir(ws, horizontal=False):
    ws.page_setup.paperSize = M.PAPEL_EN                                         # 1 = US Letter
    ws.page_setup.orientation = 'landscape' if horizontal else 'portrait'
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.oddFooter.center.text = M.FOOTER_EN
    ws.print_options.horizontalCentered = True
    ws.protection.sheet = True                                                   # sin contraseña (D20)
    ws.sheet_view.showGridLines = False


def _dv(ws, tipo, celdas, f1=None, f2=None, op=None, titulo='Invalid value', msg=''):
    dv = DataValidation(type=tipo, formula1=f1, formula2=f2, operator=op, allow_blank=False, showErrorMessage=True,
                        errorTitle=titulo, error=msg)
    ws.add_data_validation(dv)
    for c in celdas:
        dv.add(c)


# ---------------------------------------------------------------------------------------------------------------------
# Hoja "Cost per Drink"
# ---------------------------------------------------------------------------------------------------------------------
def hoja_escandallo(wb):
    ws = wb.active
    ws.title = HE
    for col, w in (('A', 46), ('B', 13), ('C', 9), ('D', 12), ('E', 12)):
        ws.column_dimensions[col].width = w
    for col in BCOL:
        ws.column_dimensions[col].width = 13
    ws.column_dimensions[OCOL].width = 58
    _f(ws, 'A1', 'COST PER DRINK — 16 INGREDIENTS × 10 DRINKS, IN GRAMS, FLUID OUNCES AND UNITS', b=True, sz=14)
    _f(ws, 'A2', 'Type only in the GREEN cells (prices, waste, units per price, quantities, mix and menu prices). Amounts are '
                 'before sales tax. The prices are kit estimates: use your roaster\'s and suppliers\'. The cost of one drink on '
                 'this tab is NOT the cost of goods in the financial plan: that one is the average of the whole mix with waste, '
                 'not a recipe.', wrap=True)
    ws.merge_cells('A2:%s2' % OCOL)
    ws.row_dimensions[2].height = 34
    for col, t in zip('ABF', ('SETTINGS', 'Value', 'Notes')):
        _f(ws, '%s4' % col, t, b=True, fill=GRIS)
    for k, rot, fmt, val, nota in ESC_IN:
        f = E[k]
        _f(ws, 'A%d' % f, rot)
        _f(ws, 'B%d' % f, val, fmt=fmt, fill=VERDE, locked=False)
        _nota(ws, f, nota, desde='F', hasta=OCOL, chars=150.0)
    for col, t in (('A', 'INGREDIENT (recipe unit · price unit)'), ('B', 'Price before sales tax'), ('C', 'Waste'),
                   ('D', 'Recipe units per price unit'), ('E', 'Net cost per recipe unit')):
        _f(ws, '%s%d' % (col, E_CAB), t, b=True, fill=GRIS, wrap=True)
    _f(ws, '%s%d' % (OCOL, E_CAB), 'WHERE THE EXAMPLE PRICE COMES FROM', b=True, fill=GRIS, wrap=True)
    _f(ws, 'A%d' % E_PVP, 'Menu price per drink', b=True)
    _f(ws, 'A%d' % E_MIX, 'Sales mix (% of drinks)', b=True)
    _f(ws, 'A%d' % E_DOB, 'Double shots per drink', b=True)
    for j, (nom, (_n, mez, dob, _leche, _oz, pvp)) in enumerate(BEB.items()):
        c = BCOL[j]
        _f(ws, '%s%d' % (c, E_CAB), nom, b=True, fill=VERDE, wrap=True, locked=False)
        _f(ws, '%s%d' % (c, E_PVP), pvp, fmt=USD, fill=VERDE, locked=False, b=True)
        _f(ws, '%s%d' % (c, E_MIX), mez / 100.0, fmt=PCT, fill=VERDE, locked=False)
        _f(ws, '%s%d' % (c, E_DOB), float(dob), fmt='0.0', fill=VERDE, locked=False)
    ws.row_dimensions[E_CAB].height = 44
    g = grid()
    for i, (nom, p, m, u, origen) in enumerate(INSUMOS):
        f = E_I0 + i
        _f(ws, 'A%d' % f, nom, fill=VERDE, locked=False)
        _f(ws, 'B%d' % f, p, fmt=USD, fill=VERDE, locked=False)
        _f(ws, 'C%d' % f, m, fmt='0%', fill=VERDE, locked=False)
        _f(ws, 'D%d' % f, u, fmt='#,##0.0', fill=VERDE, locked=False)
        _f(ws, 'E%d' % f, '=B%d/(1-C%d)/D%d' % (f, f, f), fmt='#,##0.0000')
        for j in range(NB):
            c = BCOL[j]
            if i == 0:
                _f(ws, '%s%d' % (c, f), '=%s$%d*$B$%d' % (c, E_DOB, E['gdoble']), fmt='#,##0.0;-#,##0.0;"-"')
            else:
                _f(ws, '%s%d' % (c, f), g[i][j], fmt='#,##0.00;-#,##0.00;"-"', fill=VERDE, locked=False)
        _f(ws, '%s%d' % (OCOL, f), origen, i=True, wrap=True)
    _f(ws, 'A%d' % E_R, 'RESULT PER DRINK', b=True, fill=GRIS)
    for j in range(NB):
        _f(ws, '%s%d' % (BCOL[j], E_R), '=%s%d' % (BCOL[j], E_CAB), b=True, fill=GRIS, wrap=True)
    ws.row_dimensions[E_R].height = 32
    for k, rot, fmt in RES:
        f = RW[k]
        neg = k in ('c_tot', 'fc')
        _f(ws, 'A%d' % f, rot, b=neg)
        for j in range(NB):
            c = BCOL[j]
            col_prod = '%s$%d:%s$%d' % (c, E_I0, c, E_P1)
            col_env = '%s$%d:%s$%d' % (c, E_P1 + 1, c, E_I1)
            ex = {
                'leche': 'SUM(%s%d:%s%d)' % (c, E_I0 + 1, c, E_I0 + 4),
                'c_prod': 'SUMPRODUCT(%s,$E$%d:$E$%d)' % (col_prod, E_I0, E_P1),
                'c_env': 'SUMPRODUCT(%s,$E$%d:$E$%d)' % (col_env, E_P1 + 1, E_I1),
                'c_tot': '%s%d+%s%d' % (c, RW['c_prod'], c, RW['c_env']),
                'pvp_sin': '%s%d/(1+$B$%d)' % (c, E_PVP, E['tax_menu']),
                'fc_sin': '%s%d/%s%d' % (c, RW['c_prod'], c, RW['pvp_sin']),
                'fc': '%s%d/%s%d' % (c, RW['c_tot'], c, RW['pvp_sin']),
                'sem': 'IF(%s%d<=$B$%d,"%s",IF(%s%d<=$B$%d+$B$%d,"%s","%s"))' % (
                    c, RW['fc'], E['obj'], FC_OK, c, RW['fc'], E['obj'], E['tol'], FC_VIG, FC_ALTO),
                'margen': '%s%d-%s%d' % (c, RW['pvp_sin'], c, RW['c_tot']),
                'pvp_obj': '%s%d/$B$%d*(1+$B$%d)' % (c, RW['c_tot'], E['obj'], E['tax_menu']),
            }[k]
            _f(ws, '%s%d' % (c, f), '=' + ex, fmt=fmt, b=neg, wrap=(k == 'sem'))
    ws.row_dimensions[RW['sem']].height = 32
    rng = '%s%d:%s%d' % (BCOL[0], RW['sem'], BCOL[-1], RW['sem'])
    top = '%s%d' % (BCOL[0], RW['sem'])
    for txt, fill in ((FC_OK, OK), (FC_VIG, AMBAR), (FC_ALTO, ROJO)):
        ws.conditional_formatting.add(rng, FormulaRule(formula=['%s="%s"' % (top, txt)], fill=fill))
    _f(ws, 'A%d' % E_M0, 'COST OF THE MIX (read by the "%s" tab)' % H1, b=True, fill=GRIS)
    _f(ws, 'B%d' % E_M0, 'Value', b=True, fill=GRIS)
    mix = '$%s$%d:$%s$%d' % (BCOL[0], E_MIX, BCOL[-1], E_MIX)
    fil = lambda k: '$%s$%d:$%s$%d' % (BCOL[0], RW[k], BCOL[-1], RW[k])          # noqa: E731
    suma = '$B$%d' % MR['suma']
    # VALUE() solo fuerza un número nativo en pycel (SUMPRODUCT devuelve un escalar de numpy); en Excel no cambia nada
    pond = lambda rango: 'IF(%s=0,0,VALUE(SUMPRODUCT(%s,%s))/%s)' % (suma, mix, rango, suma)      # noqa: E731
    cafe = '$%s$%d:$%s$%d' % (BCOL[0], E_I0, BCOL[-1], E_I0)
    ex = {
        'suma': 'SUM(%s)' % mix,
        'chk': 'IF(ROUND(%s,2)=1,"%s","%s")' % (suma, MZ_OK, MZ_NO),
        'm_prod': pond(fil('c_prod')), 'm_env': pond(fil('c_env')),
        'm_tot': '$B$%d+$B$%d' % (MR['m_prod'], MR['m_env']),
        'm_cafe': pond(cafe), 'm_leche': pond(fil('leche')), 'm_ticket': pond(fil('pvp_sin')),
        'm_fc': 'IF($B$%d=0,0,$B$%d/$B$%d)' % (MR['m_ticket'], MR['m_tot'], MR['m_ticket']),
    }
    for k, rot, fmt in MZ:
        f = MR[k]
        neg = k in ('m_tot', 'chk')
        _f(ws, 'A%d' % f, rot, b=neg)
        _f(ws, 'B%d' % f, '=' + ex[k], fmt=fmt, b=neg, wrap=(k == 'chk'))
    ws.merge_cells('B%d:E%d' % (MR['chk'], MR['chk']))
    ws.conditional_formatting.add('B%d' % MR['chk'], FormulaRule(formula=['$B$%d="%s"' % (MR['chk'], MZ_NO)], fill=ROJO))
    f = E_FIN
    _f(ws, 'A%d' % f, 'The all-in cost is the drink served in a single-use cup; with reusable cups only the ingredient cost '
                      'counts. Milk is a TCS food: keep it at 41 °F (5 °C) or below (the FDA Food Code cold-holding limit). The cost of '
                      'goods in the financial plan is something else: the average of the mix with waste.', i=True, wrap=True)
    ws.merge_cells('A%d:%s%d' % (f, OCOL, f))
    ws.row_dimensions[f].height = 34
    _dv(ws, 'decimal', ['B%d' % E['tax_menu'], 'B%d' % E['tol']] + ['C%d' % (E_I0 + i) for i in range(NI)],
        f1='0', f2='0.9', op='between', titulo='Invalid percentage', msg='Enter a percentage from 0 to 0.9 (0.1 = 10%).')
    _dv(ws, 'decimal', ['B%d' % E['obj']], f1='0.01', f2='0.9', op='between', titulo='Invalid target',
        msg='The target cannot be 0 (the sheet divides by it): enter 0.01 to 0.9.')
    _dv(ws, 'decimal', ['%s%d' % (c, E_MIX) for c in BCOL], f1='0', f2='1', op='between', titulo='Invalid percentage',
        msg='Enter a percentage from 0 to 1. The whole mix must add up to 100%.')
    _dv(ws, 'decimal', ['D%d' % (E_I0 + i) for i in range(NI)], f1='0.001', op='greaterThanOrEqual',
        titulo='Invalid number', msg='Enter a number above 0 (the sheet divides by it).')
    _dv(ws, 'decimal', ['B%d' % E['gdoble']] + ['B%d' % (E_I0 + i) for i in range(NI)] + ['%s%d' % (c, E_PVP) for c in BCOL]
        + ['%s%d' % (c, E_DOB) for c in BCOL] + ['%s%d' % (c, E_I0 + i) for c in BCOL for i in range(1, NI)],
        f1='0', op='greaterThanOrEqual', titulo='Invalid amount', msg='Enter a number of 0 or more.')
    ws.freeze_panes = 'B%d' % (E_CAB + 1)
    _imprimir(ws, horizontal=True)


# ---------------------------------------------------------------------------------------------------------------------
# Hoja "Pricing Calculator"
# ---------------------------------------------------------------------------------------------------------------------
def hoja_calculadora(wb):
    ws = wb.create_sheet(H1)
    for col, w in (('A', 66), ('B', 24), ('C', 17), ('D', 17), ('E', 17), ('F', 56)):
        ws.column_dimensions[col].width = w
    _f(ws, 'A1', 'COFFEE CART PRICING CALCULATOR — FROM THE COST PER DRINK TO THE QUOTE', b=True, sz=14)
    _f(ws, 'A2', 'Type only in the GREEN cells: everything else is calculated. Amounts are before sales tax unless the row says '
                 'otherwise. The defaults are kit estimates (US ratios for drinks and staffing): replace them with your own.',
       wrap=True)
    ws.merge_cells('A2:F2')
    ws.row_dimensions[2].height = 32
    for col, t in zip('ABC', ('EVENT DETAILS', 'Value', 'Notes')):
        _f(ws, '%s4' % col, t, b=True, fill=GRIS)
    for k, rot, fmt, val, nota in ENTRADAS:
        f = K[k]
        _f(ws, 'A%d' % f, rot)
        _f(ws, 'B%d' % f, val, fmt=fmt if fmt != '@' else None, fill=VERDE, locked=False, wrap=(fmt == '@'))
        _nota(ws, f, nota)
    for col, t in zip('ABC', ('CALCULATION', 'Value', 'Notes')):
        _f(ws, '%s%d' % (col, F_CALC), t, b=True, fill=GRIS)
    ref = refs('B')
    fila = F_CALC
    for k, rot, fmt, expr, nota in SALIDAS:
        fila += 1
        if expr is None:
            _f(ws, 'A%d' % fila, rot, b=True, fill=GRIS)
            continue
        neg = k in ('price', 'price_iva', 'total', 'margin_calc', 'aviso')
        _f(ws, 'A%d' % fila, rot, b=neg, wrap=len(rot) > 70)
        _f(ws, 'B%d' % fila, '=' + a_excel(expr, ref), fmt=fmt, b=neg, wrap=(fmt == 'General'))
        if fmt == 'General':
            ws.merge_cells('B%d:F%d' % (fila, fila))
            ws.row_dimensions[fila].height = 32 if k == 'aviso' else 16
            if nota:
                _f(ws, 'A%d' % fila, rot + ' · ' + nota, b=neg, wrap=True)
                ws.row_dimensions[fila].height = 48 if k == 'aviso' else 32
        else:
            _nota(ws, fila, nota)
            if len(rot) > 70:
                ws.row_dimensions[fila].height = max(ws.row_dimensions[fila].height or 16, 32)
    _f(ws, 'A%d' % F_REGLAS, 'SCALING RULES (kit estimates from US ratios)', b=True, fill=GRIS)
    _f(ws, 'B%d' % F_REGLAS, 'Value', b=True, fill=GRIS)
    for k, rot, fm, val, nota in DE.REGLAS:
        fmt = {'ent': ENT, 'dec': '#,##0.00', 'pct': PCT}[fm]
        _f(ws, 'A%d' % RG[k], rot)
        _f(ws, 'B%d' % RG[k], val, fmt=fmt, fill=VERDE, locked=False)
        _nota(ws, RG[k], nota)
    for k, rot, val in DE.CONVERSIONES:
        _f(ws, 'A%d' % CV[k], rot + ' (fixed conversion)')
        _f(ws, 'B%d' % CV[k], val, fmt='#,##0.0', fill=GRIS)
    _f(ws, 'A%d' % F_FORM, 'SERVICE LENGTH → DRINKS PER GUEST', b=True, fill=GRIS)
    _f(ws, 'B%d' % F_FORM, 'Drinks', b=True, fill=GRIS)
    for i, (nombre, hasta, v, rango) in enumerate(DE.BEBIDAS_POR_INVITADO, 1):
        _f(ws, 'A%d' % (F_FORM + i), nombre)
        _f(ws, 'B%d' % (F_FORM + i), v, fmt='0.00', fill=VERDE, locked=False)
        _f(ws, 'C%d' % (F_FORM + i), 'Reference range: %s' % rango, i=True)
    _f(ws, 'D%d' % (F_FORM + 1), 'grindkit.app ranges; %s' % DE.ORIGEN, i=True, wrap=True)
    ws.merge_cells('D%d:F%d' % (F_FORM + 1, F_FORM + 2))
    _f(ws, 'A%d' % F_HOR, 'TIME OF DAY → ADJUSTMENT TO THE DRINKS PER GUEST', b=True, fill=GRIS)
    _f(ws, 'B%d' % F_HOR, 'Adjustment', b=True, fill=GRIS)
    for i, (nombre, aj) in enumerate(HORARIOS.items(), 1):
        _f(ws, 'A%d' % (F_HOR + i), nombre)
        _f(ws, 'B%d' % (F_HOR + i), aj, fmt='0%', fill=VERDE, locked=False)
    _f(ws, 'A%d' % F_SERV, 'SERVICE OPTIONS (list for the "Served in" input)', b=True, fill=GRIS)
    for i, t in enumerate(SERVICIOS, 1):
        _f(ws, 'A%d' % (F_SERV + i), t)
    _f(ws, 'A%d' % F_ELEC, 'POWER OPTIONS (list for the "Power at the venue" input)', b=True, fill=GRIS)
    for i, t in enumerate(ELEC, 1):
        _f(ws, 'A%d' % (F_ELEC + i), t)
    _f(ws, 'A%d' % F_CAN, 'POLICY BY CLIENT TYPE', b=True, fill=GRIS)
    for col, c in zip('BCD', CANALES):
        _f(ws, '%s%d' % (col, F_CAN), c, b=True, fill=GRIS, wrap=True)
    ws.row_dimensions[F_CAN].height = 32
    for k, (rot, fmt, vals, nota) in CANAL_TAB.items():
        f = CT[k]
        _f(ws, 'A%d' % f, rot)
        for col, v in zip('BCD', vals):
            _f(ws, '%s%d' % (col, f), v, fmt=fmt, fill=VERDE, locked=False)
        _f(ws, 'E%d' % f, nota, wrap=True)
        ws.merge_cells('E%d:F%d' % (f, f))
        ws.row_dimensions[f].height = max(16, 15 * math.ceil(len(nota) / 62.0))
    for n, (k, rot) in enumerate(LABELS.items(), 1):
        f = CT[k]
        _f(ws, 'A%d' % f, '%s of cancellation (text, built from the cutoffs)' % rot)
        for col in 'BCD':
            d = lambda kk: '%s%d' % (col, CT[kk])                                 # noqa: E731
            ex = {1: '%s&" days or more"' % d('d1'),
                  2: '"From "&%s&" to less than "&%s&" days"' % (d('d2'), d('d1')),
                  3: '"From "&%s&" to less than "&%s&" days"' % (d('d3'), d('d2')),
                  4: '"Less than "&%s&" days"' % d('d3')}[n]
            _f(ws, '%s%d' % (col, f), '=' + ex, wrap=True)
        ws.row_dimensions[f].height = 32
    _f(ws, 'A%d' % F_ESC, 'SCALING TABLE (typical event: a 2-hour bar in the afternoon, single-use cups)', b=True, fill=GRIS)
    for col, n in zip('BCDE', DE.TAM_REF):
        _f(ws, '%s%d' % (col, F_ESC), n, b=True, fill=GRIS, fmt='0" guests"')
    _f(ws, 'A%d' % ESC_HORAS, 'Bar hours of the typical event')
    _f(ws, 'B%d' % ESC_HORAS, DE.HORAS_REF, fmt='0', fill=VERDE, locked=False)
    r_ = dict(('r_' + k, "$B$%d" % f) for k, f in RG.items())
    c_ = dict(('c_' + k, "$B$%d" % f) for k, f in CV.items())
    fr = F_FORM + 1 + [n for n, *_x in DE.BEBIDAS_POR_INVITADO].index(DE.FORMATO_REF)
    hr = F_HOR + 1 + list(HORARIOS).index(DE.HORARIO_REF)
    esc_cafe, esc_leche = esc_ref('m_cafe'), esc_ref('m_leche')
    kwm = '$B$%d' % K['kw_maq']
    F = dict((k, ESC_F0 + i) for i, (k, *_r) in enumerate(DE.FILAS_TABLA))
    for k, rot, fm in DE.FILAS_TABLA:
        _f(ws, 'A%d' % F[k], rot)
        for col in 'BCDE':
            n = '%s$%d' % (col, F_ESC)
            c = lambda kk: '%s%d' % (col, F[kk])                                 # noqa: E731
            ex = {
                'bebidas': 'ROUNDUP(ROUND(%s*$B$%d*(1+$B$%d),2),0)' % (n, fr, hr),
                'baristas': 'MAX(1,ROUNDUP(ROUND(%s/$B$%d*%s/%s,2),0))' % (c('bebidas'), ESC_HORAS, r_['r_pico'],
                                                                         r_['r_cap_barista']),
                'estaciones': 'ROUNDUP(%s/%s,0)' % (c('baristas'), r_['r_bar_maq']),
                'grupos': '%s*%s' % (c('estaciones'), r_['r_grupos_maq']),
                'molinos': '%s+%s' % (c('estaciones'), r_['r_molinos_extra']),
                'lb_cafe': 'ROUNDUP(%s*%s/%s+%s,1)' % (c('bebidas'), esc_cafe, c_['c_g_lb'], r_['r_purga']),
                'gal_leche': 'ROUNDUP(ROUND(%s*%s/%s,2),0)' % (c('bebidas'), esc_leche, c_['c_oz_gal']),
                'vasos': 'ROUNDUP(ROUND(%s*(1+%s),2),0)' % (c('bebidas'), r_['r_reserva']),
                'kw': '%s*%s' % (c('estaciones'), kwm),
                'toma': 'IF(%s<=$B$%d,$A$%d,"%s")' % (c('kw'), S['kw_enchufe'], F_ELEC + 1, TXT_TOMA_NO),
                'agua': 'ROUNDUP(ROUND(%s*%s+%s,2),0)' % (c('bebidas'), r_['r_agua_beb'], r_['r_agua_base']),
            }[k]
            _f(ws, '%s%d' % (col, F[k]), '=' + ex, fmt={'ent': ENT, 'dec': DEC, 'txt': 'General'}[fm], wrap=(fm == 'txt'))
    ws.row_dimensions[F['toma']].height = 32
    f = F_ULT
    _f(ws, 'A%d' % f, 'One barista on a 2-group machine makes 50-65 drinks an hour (flashquotes.com); a second station comes in when the peak does not '
                      'fit. A standard 110V outlet gives about 2.4 kW: a 220V 2-group machine needs a dedicated circuit or a '
                      'generator. Kit estimate from US ratios.', i=True, wrap=True)
    ws.merge_cells('A%d:F%d' % (f, f))
    ws.row_dimensions[f].height = 48
    rf = refs('B')
    _dv(ws, 'list', [rf['canal'].replace('$', '')], f1='"%s"' % ','.join(CANALES), titulo='Invalid client type',
        msg='Choose Private event, Corporate or Festival or market.')
    for k, f0, n, t in (('formato', F_FORM, len(FORMATOS), 'service length'), ('horario', F_HOR, len(HORARIOS), 'time of day'),
                        ('serv', F_SERV, len(SERVICIOS), 'service option'), ('elec', F_ELEC, len(ELEC), 'power option')):
        _dv(ws, 'list', [rf[k].replace('$', '')], f1='=$A$%d:$A$%d' % (f0 + 1, f0 + n), titulo='Invalid %s' % t,
            msg='Choose an option from the list.')
    _dv(ws, 'whole', ['B%d' % K['guests']], f1='1', f2='100000', op='between', titulo='Invalid number',
        msg='Enter a whole number of 1 or more.')
    _dv(ws, 'whole', ['B%d' % K[k] for k in ('km', 'free_mi', 'dias')]
        + ['%s%d' % (c, CT[k]) for k in ('d1', 'd2', 'd3') for c in 'BCD']
        + ['B%d' % RG[k] for k in ('cap_barista', 'bar_maq', 'grupos_maq', 'molinos_extra', 'agua_base', 'v_maquina',
                                    'v_enchufe', 'amp_enchufe')],
        f1='0', f2='100000', op='between', titulo='Invalid number', msg='Enter a whole number of 0 or more.')
    _dv(ws, 'decimal', ['B%d' % K['hours'], 'B%d' % ESC_HORAS], f1='0.5', f2='24', op='between', titulo='Invalid hours',
        msg='Enter 0.5 to 24 hours (it cannot be 0: the sheet divides by it).')
    _dv(ws, 'decimal', ['B%d' % K[k] for k in ('h_extra', 'dulces', 'rentals', 'gen_cost', 'kw_maq', 'wage', 'per_km', 'tkt',
                                                'canon_fijo', 'f_min')]
        + ['B%d' % RG[k] for k in ('pico', 'purga', 'loza_inv', 'agua_beb', 'factor_kw')]
        + ['B%d' % (F_FORM + i) for i in range(1, len(FORMATOS) + 1)],
        f1='0', op='greaterThanOrEqual', titulo='Invalid amount', msg='Enter a number of 0 or more.')
    _dv(ws, 'decimal', ['B%d' % K[k] for k in ('ss', 'iva')] + ['B%d' % RG['reserva']]
        + ['%s%d' % (c, CT[k]) for k in CANAL_TAB if k not in ('margen', 'd1', 'd2', 'd3') for c in 'BCD'],
        f1='0', f2='1', op='between', titulo='Invalid percentage', msg='Enter a percentage from 0 to 1 (0.10 = 10%).')
    _dv(ws, 'decimal', ['%s%d' % (c, CT['margen']) for c in 'BCD'], f1='0', f2='0.8', op='between',
        titulo='Invalid percentage', msg='Enter a target margin from 0 to 80% (0.50 = 50%): the sheet divides by 1 − margin '
                                         '− fee.')
    _dv(ws, 'decimal', ['B%d' % (F_HOR + i) for i in range(1, len(HORARIOS) + 1)], f1='-0.5', f2='1', op='between',
        titulo='Invalid adjustment', msg='Enter an adjustment from −50% to 100%.')
    ws.freeze_panes = 'A5'
    _imprimir(ws)


# ---------------------------------------------------------------------------------------------------------------------
# Hoja "3 Example Quotes"
# ---------------------------------------------------------------------------------------------------------------------
def hoja_ejemplos(wb):
    ws = wb.create_sheet(H2)
    ws.column_dimensions['A'].width = 66
    for col in 'BCD':
        ws.column_dimensions[col].width = 30
    _f(ws, 'A1', '3 EXAMPLE QUOTES — with the same formulas as the "%s" tab' % H1, b=True, sz=14)
    _f(ws, 'A2', 'Kit estimates to show how the calculation works, not market prices: change the inputs on "%s" to quote your '
                 'own events. The tables (margin and payments by client type, service lengths, scaling) are at the bottom of '
                 'that tab and the drink mix is on "%s"; they move these examples too.' % (H1, HE), wrap=True)
    ws.merge_cells('A2:D2')
    ws.row_dimensions[2].height = 32
    _f(ws, 'A4', 'EVENT DETAILS', b=True, fill=GRIS)
    for j, nombre in enumerate(EJEMPLOS):
        _f(ws, '%s4' % 'BCD'[j], nombre, b=True, fill=GRIS, wrap=True)
    ws.row_dimensions[4].height = 48
    for k, rot, fmt, val, nota in ENTRADAS:
        _f(ws, 'A%d' % K[k], rot)
        for j, e in enumerate(EJEMPLOS.values()):
            _f(ws, '%s%d' % ('BCD'[j], K[k]), e[k], fmt=fmt if fmt != '@' else None, wrap=(fmt == '@'))
    _f(ws, 'A%d' % F_CALC, 'RESULT', b=True, fill=GRIS)
    for j in range(3):
        _f(ws, '%s%d' % ('BCD'[j], F_CALC), list(EJEMPLOS)[j], b=True, fill=GRIS, wrap=True)
    ws.row_dimensions[F_CALC].height = 48
    fila = F_CALC
    for k, rot, fmt, expr, nota in SALIDAS:
        fila += 1
        if expr is None:
            _f(ws, 'A%d' % fila, rot, b=True, fill=GRIS)
            continue
        neg = k in ('price', 'price_iva', 'total', 'margin_calc', 'aviso')
        _f(ws, 'A%d' % fila, rot, b=neg, wrap=len(rot) > 70)
        for j in range(3):
            col = 'BCD'[j]
            _f(ws, '%s%d' % (col, fila), '=' + a_excel(expr, refs(col)), fmt=fmt, b=neg, wrap=(fmt == 'General'))
        if fmt == 'General':
            ws.row_dimensions[fila].height = 64 if k in ('aviso', 'tramo') else 32
        elif len(rot) > 70:
            ws.row_dimensions[fila].height = 32
    f = fila + 2
    _f(ws, 'A%d' % f, 'WHAT EACH EXAMPLE SHOWS', b=True, fill=GRIS)
    for j, (nombre, nota) in enumerate(EJ_NOTAS.items()):
        rr = calcular(EJEMPLOS[nombre])
        nota = nota.replace('{res}', '{:,.0f}'.format(rr['f_result'])).replace('{eq}', str(rr['f_eq']))
        _f(ws, 'A%d' % (f + 1 + j), nombre, b=True, wrap=True)
        _f(ws, 'B%d' % (f + 1 + j), nota, wrap=True)
        ws.merge_cells('B%d:D%d' % (f + 1 + j, f + 1 + j))
        ws.row_dimensions[f + 1 + j].height = 15 * max(3, math.ceil(len(nota) / 85.0))
    ws.freeze_panes = 'B5'
    _imprimir(ws, horizontal=True)


# ---------------------------------------------------------------------------------------------------------------------
# Hoja "Instructions"
# ---------------------------------------------------------------------------------------------------------------------
INSTRUCCIONES = [
    ('HOW TO USE — Coffee Cart Pricing Calculator (Cost per Drink to Quote per Guest)', True),
    ('Turns the cost of your drinks into a price per guest and a full quote: baristas, machines, grinders, pounds of coffee, '
     'gallons of milk, cups, power and water, subtotal, sales tax, total, payments and what you keep if the client cancels. '
     'Type only in the GREEN cells; the margin is calculated.', False),
    ('"COST PER DRINK" TAB', True),
    ('1. Enter the price before sales tax, the waste and the recipe units per price unit of each ingredient (coffee, milks, '
     'cocoa, chai, syrup, tea, sugar, cups, lid, sleeve, napkin and stirrer), and the grams, fluid ounces or units of each '
     'drink. Coffee comes from the double shots per drink and the grams per double.', False),
    ('2. For each drink you see the ingredient cost, the all-in cost with the cup, the price before sales tax, the cost of '
     'goods with a traffic light, the margin and the menu price you would need to hit your target. The sales mix must add up '
     'to 100%: the sheet warns you.', False),
    ('3. The cost of ONE drink on this tab is not the cost of goods in the financial plan: that one is the average of the mix '
     'with waste (purging, dialing in and leftover milk), not a recipe. Do not copy one into the other.', False),
    ('"PRICING CALCULATOR" TAB', True),
    ('4. Choose the client type (Private event, Corporate, or Festival or market), the guests, the service length (drinks per '
     'guest by the hours), the time of day and single-use or reusable cups. At a festival, guests are the drinks you expect '
     'to sell.', False),
    ('5. The scaling table at the bottom gives baristas, 2-group machines, grinders, pounds of coffee, gallons of milk, cups, '
     'power, the circuit you need and water. Its rules use the same assumptions as the business plan: drinks per guest by '
     'service length, 50-65 drinks per barista an hour and 18 g per double shot.', False),
    ('6. Enter the service hours and the hours outside service (loading, travel, setup and cleanup), the hourly wage and '
     'the payroll taxes. If you work the bar yourself, count your hours: it is your opportunity cost.', False),
    ('7. The price per guest is the direct cost divided by (1 − margin − fee), rounded up to the next $0.50: the lowest price '
     'that reaches your target. You also see the same price per service hour and per drink, to compare with hourly or '
     'per-drink rates.', False),
    ('8. US quotes show the price before sales tax and add the tax on the invoice. The travel fee goes on its own line, round '
     'trip beyond your free radius and with no margin. In the UK, prices to consumers include VAT: if you are VAT-registered, '
     'enter 0.20 as the tax rate.', False),
    ('POWER', True),
    ('A 2-group machine at 220V draws about 3.4 kW (NEMA 6-20 plug, 15-17 A) and needs a dedicated circuit. A standard 110-120V '
     'outlet on a 20-amp circuit gives about 2.4 kW: enough for a compact machine, not for a 220V 2-group. Grinders, fridge and '
     'lights need their own margin. Without a circuit, bring an inverter generator (5.5 kW for one machine) and ask the venue '
     'and the fire marshal where it can go. The sheet warns you, and the kit\'s contract asks for the circuit in writing.',
     False),
    ('SALES TAX, CUPS AND TIPS', True),
    ('Many states tax a hosted coffee service as catering, often on the whole invoice; at festivals and markets you collect '
     'sales tax on each drink and may need a permit where you sell. Rates and rules vary by state and city: look up yours. '
     'Some cities require compostable single-use items (Seattle, as an example). A voluntary tip or a tip jar belongs to the '
     'team: it stays out of the price, the taxable amount and the total; put out a tip jar only if the host agrees.', False),
    ('RETAINER, BALANCE AND CANCELLATION', True),
    ('Private events: a 30% retainer at booking (credited toward the price; US operators ask 30-50%) and the 70% balance 14 '
     'days before the event, with the final guest count 10 days before. Corporate: 30% at signing and 70% net 30 after the '
     'invoice. Festivals and markets: paid at the point of sale. Send a receipt for every payment.', False),
    ('Cancellation, on the total, in tiers from X to less than Y with no gaps. Private events: 60 days or more, you keep 10%; '
     'from 30 to less than 60 days, 30%; from 14 to less than 30 days, 60%; less than 14 days, 100% minus the costs you save '
     '(coffee, milk and cups you do not use, staff and travel). Corporate: 30 days or more, 10%; from 15 to less than 30 days, '
     '30%; from 7 to less than 15 days, 60%; less than 7 days, as for private events. It works both ways: if you cancel, you '
     'refund what you collected and pay the same tier. If you change a cutoff in the table, the tiers rewrite themselves. '
     'Have an attorney in your state review your contract.', False),
    ('FESTIVALS AND FARMERS MARKETS', True),
    ('Selling drink by drink needs the organizer\'s vendor terms and, usually, a temporary food event permit. The sheet '
     'calculates the sales (drinks × average price, or the organizer\'s minimum sales guarantee if it is higher), the fee (% '
     'of sales plus any flat booth fee), the result of the day and the drinks you need to cover your costs. The "price per '
     'guest" in that column is the lowest price per drink that covers cost, fee and margin.', False),
    ('FOOD SAFETY', True),
    ('Milk is a TCS food: keep it at 41 °F (5 °C) or below in the cart\'s fridge and in coolers (the FDA Food Code cold-holding limit). '
     'Use a separate pitcher for each milk, wash your hands at the cart\'s hand sink first, and buy pastries only from a '
     'licensed bakery.', False),
    ('KIT ESTIMATES', True),
    ('The defaults and the three examples are kit estimates to show the calculation, not market prices. Drinks per guest, '
     'drinks per barista and water use come from US ratios: replace them with your own measurements. Use your roaster\'s and '
     'suppliers\' prices, your wages and your miles before you quote.', False),
    ('Sources: AI Chef Pro analysis (October 2026) of flashquotes.com booking data, the grindkit.app pricing guide, BLS average '
     'prices, coffeecartboys.com and the FDA Food Code.', False),
]


def hoja_instrucciones(wb):
    ws = wb.create_sheet(H3)
    ws.column_dimensions['A'].width = 110
    f = 1
    for t, b in INSTRUCCIONES:
        _f(ws, 'A%d' % f, t, b=b, sz=13 if f == 1 else 11, wrap=True)
        if not b:
            ws.row_dimensions[f].height = max(16, 15 * math.ceil(len(t) / 120.0))
        f += 2 if b and f > 1 else 1
    f += 1
    for t in ('To edit a cell that is not green: Review → Unprotect Sheet (there is no password).',
              'More templates and kits in the catalog at ' + M.URL_TIENDA, M.FIRMA_EN, M.version_en('cof'), M.AVISO_EN):
        _f(ws, 'A%d' % f, t, wrap=True, i=t == M.AVISO_EN)
        f += 1
    _imprimir(ws)


def docprops(wb):
    p = M.PLANES['cof']
    pr = wb.properties
    pr.title = '%s · %s' % (TITULO, p['producto'])
    pr.subject = '%s · v%s' % (p['producto'], M.VERSION)
    pr.keywords = p['keywords']
    pr.description = p['url']
    pr.category = 'AI Chef Pro · Digital products'
    pr.creator = 'AI Chef Pro'
    pr.lastModifiedBy = 'AI Chef Pro'


def generar(destino):
    wb = openpyxl.Workbook()
    hoja_escandallo(wb)
    hoja_calculadora(wb)
    hoja_ejemplos(wb)
    hoja_instrucciones(wb)
    docprops(wb)
    os.makedirs(os.path.dirname(destino), exist_ok=True)
    wb.save(destino)
    sys.path.insert(0, PD)
    import inject_cache as IC                                                    # noqa: E402
    nf, inj, fail = IC.inject(destino)
    print('%s: formulas=%d cache_inyectado=%d fallos_pycel=%d' % (os.path.basename(destino), nf, inj, fail))
    return fail


# ---------------------------------------------------------------------------------------------------------------------
# Autotest (disciplina del ES + gates EN de la paella)
# ---------------------------------------------------------------------------------------------------------------------
PROHIBIDOS = [r'\bfreelance', r'\bresearch\b', r'\bSPEC\b', r'§', r'escalado_cof|datos_escalado', r'\bnon-?refundable\b',
              r'\bguarantee[ds]? (?:profit|returns?|income)', r'growing demand', r'espresso martini|liquor|alcohol',
              r'\bper day\b', r'\bcustomers\b', r'\bIVA\b', r'€', r'\bEUR\b', r'\bkg\b(?![^()]*\))', r'\bliters?\b(?![^()]*\))']


def _distinto(x, v):
    if isinstance(v, str):
        return x != v
    return not isinstance(x, (int, float)) or abs(float(x) - v) > 0.005


def celdas_escandallo(esc=None):
    esc = esc or ESC
    out = []
    for i in range(NI):
        out.append(('E%d' % (E_I0 + i), esc['neto'][i], 'neto%d' % i))
    for j in range(NB):
        for k in RW:
            out.append(('%s%d' % (BCOL[j], RW[k]), esc['bebidas'][j][k], '%s[%d]' % (k, j)))
    for k in MR:
        out.append(('B%d' % MR[k], esc['mz'][k], k))
    return out


def autotest(destino):
    from pycel import ExcelCompiler
    err, n = [], 0
    xl = ExcelCompiler(filename=destino)
    wv = openpyxl.load_workbook(destino, data_only=True)
    wf = openpyxl.load_workbook(destino)
    for c, v, lab in celdas_escandallo():
        n += 1
        for et, x in (('pycel', xl.evaluate("'%s'!%s" % (HE, c))), ('caché', wv[HE][c].value)):
            if _distinto(x, v):
                err.append('%s %s (%s): %s %r ≠ esperado %r' % (HE, c, lab, et, x, v))
    casos = [(H2, 'BCD'[j], e, nombre) for j, (nombre, e) in enumerate(EJEMPLOS.items())]
    casos.append((H1, 'B', OrderedDict(_BASE), 'hoja principal'))
    for hoja, col, e, nombre in casos:
        esperado = calcular(e)
        for k, v in esperado.items():
            c = '%s%d' % (col, S[k])
            n += 1
            for lab, x in (('pycel', xl.evaluate("'%s'!%s" % (hoja, c))), ('caché', wv[hoja][c].value)):
                if _distinto(x, v):
                    err.append('%s · %s %s (%s): %s %r ≠ esperado %r' % (nombre, hoja, c, k, lab, x, v))
        if esperado['margin_calc'] + 1e-9 < esperado['margen']:
            err.append('%s: margen calculado %.4f < objetivo %.4f' % (nombre, esperado['margin_calc'], esperado['margen']))
    # ejemplos dentro de su línea de la mezcla del plan (SPEC D227: ±10 % de $15 / $12); festival: rentable a $6.50
    for nombre, obj in OBJETIVO_EJ.items():
        r = calcular(EJEMPLOS[nombre])
        if abs(r['price'] / obj - 1) > 0.10 + 1e-9:
            err.append('%s: $%.2f por invitado fuera de ±10 %% de $%.2f' % (nombre, r['price'], obj))
    rf = calcular(EJEMPLOS['Community festival, guest-pay'])
    if not (rf['price'] <= EJEMPLOS['Community festival, guest-pay']['tkt'] and rf['f_result'] > 0 and rf['f_eq'] > 0):
        err.append('festival: precio mínimo %.2f, resultado %.2f, equilibrio %d' % (rf['price'], rf['f_result'], rf['f_eq']))
    if abs(EJEMPLOS['Community festival, guest-pay']['tkt'] / 6.50 - 1) > 0.10:
        err.append('festival: precio por bebida fuera de ±10 % de $6.50')
    # tabla de escalado (la de escalado_cof, la del Word) contra pycel
    F = dict((k, ESC_F0 + i) for i, (k, *_r) in enumerate(DE.FILAS_TABLA))
    for j, col in enumerate('BCDE'):
        d = DE.tabla_escalado()[j][1]
        for k, _rot, _fm in DE.FILAS_TABLA:
            n += 1
            x = xl.evaluate("'%s'!%s%d" % (H1, col, F[k]))
            if _distinto(x, d[k]):
                err.append('escalado %s %s: %r ≠ %r' % (col, k, x, d[k]))
    # casos propios: cambian celdas de la hoja principal y se recalculan con pycel
    base = OrderedDict(_BASE)
    propios = [('propio 1: Corporate, 40 invitados, hasta 2 h, mañana, salida estándar', dict(
        base, canal=CANALES[1], guests=40, formato='Up to 2 hours', horario='Morning (breakfast and mid-morning coffee)',
        hours=2, h_extra=4, dulces=0.0, km=0, elec=ELEC[0], kw_maq=1.8)),
        ('propio 2: Private, 260 invitados, jornada, sin salida, loza', dict(
            base, canal=CANALES[0], guests=260, formato='Over 6 hours (full day)', horario='Afternoon or mixed day',
            hours=9, elec=ELEC[2], serv=SERVICIOS[1], km=60, dias=10, rentals=150.0)),
        ('propio 3: Festival, 400 bebidas, booth fee y garantía que manda', dict(
            base, canal=CANALES[2], guests=400, hours=11, elec=ELEC[0], km=20, canon_fijo=120.0, tkt=3.2, f_min=2000.0,
            dias=3, dulces=0.2))]
    for nombre, e in propios:
        esperado = calcular(e)
        xl2 = ExcelCompiler(filename=destino)
        for k in S:
            xl2.evaluate("'%s'!B%d" % (H1, S[k]))
        for k, v in e.items():
            xl2.set_value("'%s'!B%d" % (H1, K[k]), v)
        for k, v in esperado.items():
            n += 1
            x = xl2.evaluate("'%s'!B%d" % (H1, S[k]))
            if _distinto(x, v):
                err.append('%s: %s pycel %r ≠ esperado %r' % (nombre, k, x, v))
    if calcular(propios[0][1])['toma'] != ELEC[0] or calcular(propios[1][1])['c_gen'] <= 0:
        err.append('la salida estándar o el generador no cambian el cálculo')
    if calcular(propios[2][1])['f_venta'] != 2000.0:
        err.append('la minimum sales guarantee no manda cuando es mayor que la venta')
    if calcular(dict(base, elec=ELEC[0]))['aviso'] != TXT_AV_MAL:
        err.append('el aviso no salta con una máquina de 220V en una salida estándar')
    e = dict(EJEMPLOS['Wedding coffee bar after dinner'])
    cs = [calcular(dict(e, canal=c)) for c in CANALES]
    if cs[0]['cb2'] == 0 or cs[1]['cb2'] != 0 or cs[1]['cb3'] == 0 or cs[2]['cb4'] == 0 or cs[0]['cb4'] != 0:
        err.append('el selector de cliente no cambia el cobro')
    if not calcular(dict(e, km=40))['desp'] > 0 or calcular(dict(e, km=15))['desp'] != 0:
        err.append('el radio gratis del travel fee no funciona')
    for ic in range(3):
        s = sum(CANAL_TAB[k][2][ic] for k in ('pc_senal', 'pc_antes', 'pc_30d', 'pc_acto'))
        if abs(s - 1.0) > 1e-9:
            err.append('cobro %s suma %.2f ≠ 100 %%' % (CANALES[ic], s))
    for ic in (0, 1):
        d = [CANAL_TAB[k][2][ic] for k in ('d1', 'd2', 'd3')]
        v = [CANAL_TAB['cn%d' % i][2][ic] for i in (1, 2, 3, 4)]
        if not (d[0] > d[1] > d[2] > 0):
            err.append('cancelación %s: cortes no descendentes %r' % (CANALES[ic], d))
        if v != sorted(v) or v[-1] != 1.0:
            err.append('cancelación %s: tramos no crecientes %r' % (CANALES[ic], v))
        for dia in range(0, 400):
            tr = calcular(dict(base, canal=CANALES[ic], dias=dia))['tramo']
            esp = etiquetas(*d)[0 if dia >= d[0] else (1 if dia >= d[1] else (2 if dia >= d[2] else 3))]
            if tr != esp:
                err.append('cancelación %s: el día %d cae en %r, no en %r' % (CANALES[ic], dia, tr, esp))
                break
    if [CANAL_TAB['cn%d' % i][2][ic] for i in (1, 2, 3, 4) for ic in (0, 1)] != [0.10, 0.10, 0.30, 0.30, 0.60, 0.60, 1.0, 1.0]:
        err.append('cancelación: los porcentajes no son 10 / 30 / 60 / 100 (política del ES, D16)')
    if CANAL_TAB['d3'][2][0] != 14 or CANAL_TAB['pc_senal'][2][0] != 0.30:
        err.append('política privada: tercer corte = 14 días (saldo) y retainer 30 % (textos EN del checklist)')
    if sum(b[1] for b in DE.BEBIDAS) != 100 or round(ESC['mz']['suma'], 9) != 1.0:
        err.append('la mezcla de bebidas no suma el cien por cien: %r' % ESC['mz']['suma'])
    if abs(sum(s for _n, s in DE.LECHES) - 1.0) > 1e-9:
        err.append('el reparto de leches no suma 1')
    for j, (nom, b) in enumerate(BEB.items()):
        if abs(sum(grid()[i][j] for i in range(1, 5)) - b[3]) > 1e-9:
            err.append('la leche de %s no suma %s fl oz' % (nom, b[3]))
    if abs(ESC['mz']['m_cafe'] - DE.cafe_g_medio()) > 1e-9 or abs(ESC['mz']['m_leche'] - DE.leche_oz_medio()) > 1e-9:
        err.append('café/leche medios de "Cost per Drink" ≠ escalado_cof')
    # coste medio por bebida en la banda de grindkit.app ($0.75-1.50 todo incluido) y del caso (ingredientes ≈ $1.03)
    if not 0.75 <= ESC['mz']['m_tot'] <= 1.50 or not 0.85 <= ESC['mz']['m_prod'] <= 1.20 or not 0.20 <= ESC['mz']['m_env'] <= 0.40:
        err.append('coste medio por bebida fuera de banda: %.3f + %.3f' % (ESC['mz']['m_prod'], ESC['mz']['m_env']))
    if not 0.20 <= ESC['mz']['m_fc'] <= 0.30:
        err.append('COGS de la mezcla fuera de 20-30 %%: %.3f' % ESC['mz']['m_fc'])
    alto = calcular_escandallo(pvp=[0.5] * NB)
    if any(b['sem'] != FC_ALTO for b in alto['bebidas']) or any(
            b['sem'] != FC_OK for b in calcular_escandallo(pvp=[40.0] * NB)['bebidas']):
        err.append('el semáforo no distingue los estados')
    ws = wf[H1]
    dvs = dict((str(rng), dv) for dv in ws.data_validations.dataValidation for rng in dv.sqref.ranges)
    c = 'B%d' % K['canal']
    if dvs.get(c) is None or dvs[c].type != 'list' or ','.join(CANALES) not in (dvs[c].formula1 or ''):
        err.append('DV de lista incorrecta en %s' % c)
    for k in ('formato', 'horario', 'serv', 'elec'):
        dv = dvs.get('B%d' % K[k])
        if dv is None or dv.type != 'list' or not (dv.formula1 or '').startswith('=$A$'):
            err.append('DV de lista del rango incorrecta en %s' % k)
    for k in ('guests', 'km', 'free_mi', 'dias'):
        dv = dvs.get('B%d' % K[k])
        if dv is None or dv.type != 'whole':
            err.append('DV entera ausente en %s' % k)
    dve = dict((str(rng), dv) for dv in wf[HE].data_validations.dataValidation for rng in dv.sqref.ranges)
    for c in ('B%d' % E['obj'], '%s%d' % (BCOL[0], E_MIX), '%s%d' % (BCOL[0], E_I0 + 1), 'D%d' % E_I0):
        if dve.get(c) is None:
            err.append('%s: falta DV en %s' % (HE, c))
    textos = []
    for w in wf.worksheets:
        if w.page_setup.paperSize != 1 or not w.sheet_properties.pageSetUpPr.fitToPage or w.oddFooter.center.text != M.FOOTER_EN:
            err.append('%s: impresión (Letter + ajuste + pie)' % w.title)
        if not w.protection.sheet or w.protection.password:
            err.append('%s: protección sin contraseña' % w.title)
        for row in w.iter_rows():
            for c in row:
                v = c.value
                if v == '':
                    err.append('%s!%s: cadena vacía' % (w.title, c.coordinate))
                if isinstance(v, str) and c.data_type != 'f':
                    textos.append((w.title, c.coordinate, v))
                if isinstance(v, str) and c.data_type == 'f':
                    textos += [(w.title, c.coordinate + ' (fórmula)', lit) for lit in re.findall(r'"((?:[^"]|"")*)"', v)
                               if any(ch.isalpha() for ch in lit)]
                if c.data_type == 'f' and wv[w.title][c.coordinate].value is None:
                    err.append('%s!%s: fórmula sin caché' % (w.title, c.coordinate))
                if c.fill.fgColor.rgb in ('00E8F5E9', 'FFE8F5E9') and c.protection.locked:
                    err.append('%s!%s: verde bloqueada' % (w.title, c.coordinate))
                fmt = c.number_format or ''
                if '€' in fmt or '$' in fmt:
                    err.append('%s!%s: símbolo de moneda en el formato (D8)' % (w.title, c.coordinate))
        for dv in w.data_validations.dataValidation:
            for t in (dv.error, dv.errorTitle, dv.prompt, dv.promptTitle):
                if t:
                    textos.append((w.title, 'DV %s' % dv.sqref, t))
    for w in wf.worksheets:                     # ninguna tasa tecleada en una fórmula: solo 0, 1, 2 y 1000
        for row in w.iter_rows():
            for c in row:
                if c.data_type == 'f':
                    sin = re.sub(r'"[^"]*"', '', str(c.value))
                    sin = re.sub(r"'[^']+'!|\$?[A-Z]{1,2}\$?\d+", '', sin)
                    for num in re.findall(r'(?<![A-Za-z])\d+\.?\d*', sin):
                        if num not in ('0', '1', '2', '1000'):
                            err.append('%s!%s: número tecleado %s en la fórmula %s' % (w.title, c.coordinate, num, c.value))
    p = wf.properties
    for campo in ('title', 'subject', 'keywords', 'description'):
        textos.append(('docProps', campo, getattr(p, campo) or ''))
    for h, c, t in textos:
        r = [x for x in M.restos_espanol(t) if not x[0].startswith('oficio')]
        if r:
            err.append('%s!%s: restos de español %s' % (h, c, r[:1]))
        if M.no_latinos(t):
            err.append('%s!%s: caracteres no latinos' % (h, c))
        if M.cifras_prohibidas(t):
            err.append('%s!%s: cifra/término prohibido %s' % (h, c, M.cifras_prohibidas(t)[:1]))
        if re.search(r'(?i)\bvalue\b.*\$\d|\bsave \$|total value|\bwas \$|\d+\s?% off\b', t):
            err.append('%s!%s: ancla de precio' % (h, c))
        for rx in PROHIBIDOS:
            if re.search(rx, t, re.I):
                err.append('%s!%s: prohibido /%s/' % (h, c, rx))
    todo = '\n'.join(t for _h, _c, t in textos)
    if M.version_en('cof') not in todo or M.FIRMA_EN not in todo:
        err.append('falta la línea de versión o la firma')
    for frag in ('2.4 kW', '3.4 kW', 'NEMA 6-20', '41 °F'):
        if frag not in todo:
            err.append('falta "%s" (aviso eléctrico / Food Code)' % frag)
    if wf.sheetnames != [HE, H1, H2, H3]:
        err.append('pestañas %r' % wf.sheetnames)
    if p.creator != 'AI Chef Pro' or not (p.title or '').startswith(TITULO) or p.description != M.PLANES['cof']['url']:
        err.append('docProps D20: %r / %r / %r' % (p.creator, p.title, p.description))
    for x in err[:80]:
        print('ERROR', x)
    print('autotest calculadora coffee cart EN: %d comprobaciones (pycel + caché en "%s", ejemplos y hoja principal; pycel en '
          '3 casos propios y en la tabla de escalado) · %d textos · %s' % (
              n * 2, HE, len(textos), 'VERDE' if not err else 'ROJO (%d errores)' % len(err)))
    mz = ESC['mz']
    print('  %s: %.3f ingredientes + %.3f vaso = %.3f por bebida · carta media %.2f · COGS %.1f %% · café %.2f g · leche %.2f '
          'fl oz' % (HE, mz['m_prod'], mz['m_env'], mz['m_tot'], mz['m_ticket'], mz['m_fc'] * 100, mz['m_cafe'], mz['m_leche']))
    for nombre, e in EJEMPLOS.items():
        r = calcular(e)
        print('  %-34s %4d beb · %d bar · %d máq · %6.2f/inv · %6.2f con tax · %7.2f/h · %5.2f/bebida · total %9.2f · margen '
              '%.1f %%' % (nombre, r['bebidas'], r['baristas'], r['estaciones'], r['price'], r['price_iva'], r['price_h'],
                           r['price_taza'], r['total'], r['margin_calc'] * 100))
        if e['canal'] == CANALES[2]:
            print('     festival: venta %.2f · cuota %.2f · coste %.2f · resultado %.2f · equilibrio %d bebidas' % (
                r['f_venta'], r['f_canon'], r['c_total'], r['f_result'], r['f_eq']))
    print('  escalado: ' + ' · '.join('%d inv: %d beb, %d bar, %d máq, %.1f lb, %d gal leche, %s' % (
        nn, d['bebidas'], d['baristas'], d['estaciones'], d['lb_cafe'], d['gal_leche'], d['toma']) for nn, d in
        DE.tabla_escalado()))
    return 1 if err else 0


if __name__ == '__main__':
    ap = argparse.ArgumentParser()
    ap.add_argument('--salida', default=DIR_DEF, help='carpeta de salida (por defecto, dl/coffee-cart-business-plan/)')
    ap.add_argument('--autotest', action='store_true')
    a = ap.parse_args()
    destino = os.path.join(a.salida, FICHERO)
    if not a.autotest and generar(destino):
        sys.exit('ERROR: fallos de pycel en inject_cache')
    sys.exit(autotest(destino))
