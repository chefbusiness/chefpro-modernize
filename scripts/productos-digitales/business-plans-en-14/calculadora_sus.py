#!/usr/bin/env python3
"""
calculadora_sus.py — Sushi Fish, Rice & Costing Calculator, calculadora del Sushi Restaurant Business Plan Kit (EN, US).
Es la edición EE. UU. de `../plan-sushi/calculadora_sus_es.py`: MISMAS hojas (`Fish Yield` · `Sushi Rice` · `Costing` ·
`Instructions`), filas, fórmulas, helpers, `inject_cache` y autotest (pycel + caché = `calcular()`, 3 juegos), con el patrón
de `../business-plans-en-13/calculadora_pol.py`. Cambios EN:

  (1) pescado en libras ($/lb) y porciones en ONZAS, con la constante gris «16 oz per lb» (el gate prohíbe números tecleados
      en fórmulas salvo 0/1); vinagre en fl oz por lb de arroz cocido y $/galón con la constante gris «128 fl oz per gallon»;
      «60 minutes per hour»; °F solo en el texto de las instrucciones;
  (2) coste por lb útil = precio ÷ (rendimiento + recorte) ÷ (1 − merma de case); mano de obra del head sushi chef en $/h;
  (3) regla de parásitos FDA Food Code 2022 (3-402.11 / 3-402.12) en el desplegable informativo; arroz de sushi: TPHC 4 h
      (3-501.19) o variance con plan HACCP (3-502.11), sin fijar un pH legal;
  (4) hoja `Costing`: sales tax en carta = 0 %, comisión del marketplace 25 %; margen dine-in · take-out/pickup · marketplace;
  (5) US Letter, formatos sin símbolo de moneda (los rótulos llevan «($)»), docProps en inglés, versión del libro EN.

Valores que otros scripts leen de la CACHÉ del xlsx (dict `EXPUESTOS`, función `valores_expuestos(ruta)`).
Corre en el VPS (openpyxl + pycel):

    /root/venv-guias/bin/python calculadora_sus.py --salida <dir>             # genera + inject_cache + autotest
    /root/venv-guias/bin/python calculadora_sus.py --salida <dir> --autotest

Sesión Claude Code, 31-oct-2026.
"""
import argparse
import math
import os
import re
import sys
from collections import OrderedDict

import openpyxl
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill, Protection
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

AQUI = os.path.dirname(os.path.abspath(__file__))
PD = os.path.dirname(AQUI)
os.environ.setdefault('BP_CONJUNTO', 'sus')
sys.path.insert(0, os.path.join(PD, 'planes-v2_0'))
sys.path.insert(0, os.path.join(PD, 'business-plans-en'))
sys.path.insert(0, PD)
sys.path.insert(0, AQUI)
import motor                                                                       # noqa: E402  (solo gate_formatos)
import mapas as M                                                                  # noqa: E402  (detectores y textos fijos)

_PL = M.PLANES['sus']
SLUG = _PL['slug']
PRODUCTO = _PL['producto']
URL = _PL['url']
FICHERO = 'sushi-fish-rice-costing-calculator.xlsx'
TITULO = 'Sushi Fish, Rice & Costing Calculator (Fish Yield, Sushi Rice and Cost per Piece)'
KEYWORDS = _PL['keywords']
VERSION, MES = M.VERSION, M.MES_VERSION
VERSION_LINEA = 'Version %s · %s · %s · info@aichef.pro' % (VERSION, MES, URL)
FOOTER = M.FOOTER_EN
HP, HA, HE, HI = 'Fish Yield', 'Sushi Rice', 'Costing', 'Instructions'
VERDE = PatternFill('solid', fgColor='E8F5E9')
GRIS = PatternFill('solid', fgColor='F2F2F2')
ROJO = PatternFill('solid', bgColor='F8CBAD', fgColor='F8CBAD')
AMBAR = PatternFill('solid', bgColor='FFE699', fgColor='FFE699')
OK = PatternFill('solid', bgColor='C6E0B4', fgColor='C6E0B4')
FUENTE = 'Calibri'
USD, ENT, PCT, DEC, OZF, USD3 = '#,##0.00', '#,##0', '0.0%', '#,##0.0', '#,##0.00', '#,##0.000'
FC_OK, FC_VIG, FC_ALTO = 'On target', 'Watch the cost', 'Above target'
SEM_ENT, SEM_LOMO = 'The whole fish pays off', 'Buy the fillet'
EST = ' Kit estimate, editable.'

# ---------------------------------------------------------------------------------------------------------------------
# Datos (F1-research-us-delta §4-§5). Importes antes de sales tax. Todo «Kit estimate (Oct 2026)», editable.
# ---------------------------------------------------------------------------------------------------------------------
NPE = 4
PC = [get_column_letter(2 + i) for i in range(NPE)]                                # B..E
#: hoja Fish Yield, bloque A: (nombre, precio $/lb, rendimiento del lomo, recorte reaprovechado)
PESCADOS = [
    ('Whole farmed Atlantic salmon', 7.50, 0.55, 0.15),
    ('Salmon fillet, skin-on, pin-bone out', 20.00, 0.85, 0.03),
    ('Yellowfin tuna saku block (frozen)', 23.00, 0.90, 0.05),
    ('Other fish (type yours)', 24.50, 0.80, 0.05),
]
P_ROW = OrderedDict([('pr', 6), ('rd', 7), ('rc', 8)])
P_NOTAS = OrderedDict([
    ('pr', 'No wholesale price was read for sushi salmon in Oct 2026: NYC auction reports $18-24/lb for sushi salmon (weak source) and yellowfin '
           'saku retails at about $22.99/lb (Costco). The kit puts 7.50 · 20.00 · 23.00 · 24.50 $/lb. Kit estimate (Oct 2026): replace with your '
           "distributor's price."),
    ('rd', 'Usable loin is 45-60% of a whole salmon (secondary sources). Measure it with 3 fish (see the "Instructions" sheet). For a saku block '
           'it is the share you keep after trimming.' + EST),
    ('rc', 'Trim you reuse (spicy mix, maki, poke), as a % of the weight you buy. What you throw away does not count here.' + EST),
])
OPC = [  # desplegable informativo: (opción, recordatorio)
    ('I freeze in-house: −4 °F for 7 days or a −31 °F option',
     'FDA Food Code 2022, 3-402.11: −4 °F or below for 168 hours (7 days), or −31 °F until solid and 15 hours, or −31 °F until solid and 24 hours '
     'at −4 °F. Log time and temperature of every lot and keep the records 90 days (3-402.12).'),
    ('Arrives frozen with a supplier letter',
     "Keep the supplier's written agreement or statement for every lot: it must show the freezing time and temperature. Keep it 90 days after you "
     'serve the fish (3-402.12).'),
    ('Exempt: tuna species or feed-raised aquaculture',
     'Tuna (Thunnus species) is exempt; feed-raised aquaculture fish such as farmed salmon is exempt only with a written statement from the supplier or '
     'the farm. Keep the tuna invoice that names the species, file the statement and ask again if the source changes.'),
]
#: bloque B (entradas comunes de la hoja Fish Yield): (clave, rótulo, formato, valor, constante, nota)
ENT_PE = [
    ('m', 'Case shrink (% of cut fish that is not sold)', PCT, 0.05, False,
     'Fish left in the sushi case or at the pass and thrown out at close.' + EST),
    ('gn', 'Ounces of fish per nigiri slice (oz)', OZF, 0.5, False, 'Kit standard: weigh 10 slices and use the average.' + EST),
    ('gs', 'Ounces per sashimi cut (oz)', OZF, 0.5, False, 'Weigh 10 cuts and use the average.' + EST),
    ('gr', 'Ounces of fish per roll (oz)', OZF, 1.5, False, 'A roll of 6 to 8 pieces with fish inside.' + EST),
    ('tl', 'Cleaning minutes per whole fish (min)', DEC, 25.0, False,
     'No source: time it. Time the cleaning of 3 fish and use the average.' + EST),
    ('hh', 'Head sushi chef cost per hour ($/h)', USD, 33.0, False,
     '$5,199 a month × 12 + 10% payroll taxes ÷ 2,080 h ≈ $33/h.' + EST),
    ('pe', 'Weight of a whole fish (lb)', '#,##0.00', 11.0, False, 'Whole farmed salmon of about 11 lb.' + EST),
    ('cong', 'Parasite destruction: freezing or statement (dropdown)', 'General', OPC[0][0], False,
     'Informational: pick your case and read the reminder below. It changes no figure.'),
    ('OZ', 'Ounces per lb (constant: 16 oz per lb)', ENT, 16, True, 'Constant: 1 lb = 16 oz. Not editable.'),
    ('MH', 'Minutes per hour (constant: 60 minutes per hour)', ENT, 60, True, 'Constant: 1 hour = 60 minutes. Not editable.'),
]
PE_M0 = 11
M_PE = OrderedDict((k[0], PE_M0 + i) for i, k in enumerate(ENT_PE))                # 11..20
PE_RC = PE_M0 + len(ENT_PE) + 1                                                    # 22: cabecera del bloque C
SAL_C = [
    ('tot', 'Share used (yield + trim)', PCT, '{rd}+{rc}', 'Loin plus the trim you reuse.'),
    ('cku', 'COST PER USABLE LB ($)', USD, '{pr}/{tot}/(1-{m})', 'Price ÷ (yield + trim) ÷ (1 − case shrink).'),
    ('c_nig', 'Cost of a nigiri slice ($)', USD3, '{cku}*{gn}/{OZ}', 'Cost per usable lb × ounces ÷ 16.'),
    ('c_sas', 'Cost of a sashimi cut ($)', USD3, '{cku}*{gs}/{OZ}', 'Cost per usable lb × ounces ÷ 16.'),
    ('c_roll', 'Cost of the fish in a roll ($)', USD3, '{cku}*{gr}/{OZ}', 'Cost per usable lb × ounces ÷ 16.'),
]
S_C = OrderedDict((k[0], PE_RC + 1 + i) for i, k in enumerate(SAL_C))             # 23..27
PE_RD = PE_RC + len(SAL_C) + 2                                                     # 29: cabecera del bloque D
SAL_D = [
    ('lab_pz', 'Cost of cleaning one whole fish ($)', USD, '{tl}/{MH}*{hh}', 'Cleaning minutes ÷ 60 × head sushi chef cost per hour.'),
    ('kg_pz', 'Usable lb of one whole fish (lb)', '#,##0.00', '{pe}*{rd0}', 'Weight of the fish × yield of the whole salmon.'),
    ('lab_kg', 'CLEANING LABOR PER USABLE LB ($)', USD, '{lab_pz}/{kg_pz}', 'Cost of cleaning one fish ÷ usable pounds.'),
    ('c_ent', 'Cost per usable lb of the whole fish with cleaning included ($)', USD, '{cku0}+{lab_kg}',
     'Cost per usable lb + cleaning labor per usable lb.'),
    ('X', 'THE WHOLE FISH PAYS OFF AS LONG AS CLEANING COSTS LESS THAN ($ per usable lb)', USD, '{cku1}-{cku0}',
     'Cost per usable lb of the fillet − cost per usable lb of the whole fish.'),
    ('sem', 'Verdict with your labor', 'General', 'IF({lab_kg}<{X},"%s","%s")' % (SEM_ENT, SEM_LOMO),
     'The whole fish pays off if cleaning costs less than that difference.'),
    ('rem', 'Reminder for your parasite case', 'General', 'INDEX({rem_r},MATCH({cong},{opc_r},0))', 'Text from the list below.'),
]
S_D = OrderedDict((k[0], PE_RD + 1 + i) for i, k in enumerate(SAL_D))             # 30..36
PE_RO = PE_RD + len(SAL_D) + 2                                                     # 38: cabecera de la lista
PE_O0 = PE_RO + 1                                                                  # 39..41
PE_O1 = PE_O0 + len(OPC) - 1

#: hoja Sushi Rice: (clave, rótulo, formato, valor, constante, nota)
ENT_A = [
    ('ak', 'Raw sushi rice ($/lb, before tax)', USD, 1.04, False,
     '50 lb bag $51.99 = $1.04/lb, webstaurantstore.com, Oct 2026; other medium-grain brands $1.24-1.28/lb.'),
    ('f', 'Raw to cooked factor (×)', '#,##0.00', 2.2, False, '1 lb of raw rice gives 2.2 lb of cooked rice.' + EST),
    ('z', 'Sushi vinegar (sushi-zu) per lb of cooked rice (fl oz)', DEC, 1.5, False,
     'Rice vinegar, sugar and salt: about 1.5 fl oz per lb of cooked rice.' + EST),
    ('zp', 'Sushi vinegar ($/gallon, before tax)', USD, 20.50, False, 'Seasoned rice vinegar about $20.50/gal, webstaurantstore.com, Oct 2026.'),
    ('g_nig', 'Ounces of sushi rice per nigiri (oz)', OZF, 0.75, False, 'Weigh 10 rice balls and use the average.' + EST),
    ('g_hos', 'Ounces of sushi rice per maki (thin roll) (oz)', OZF, 2.5, False, 'A thin roll of 6 pieces.' + EST),
    ('g_ura', 'Ounces of sushi rice per inside-out roll (oz)', OZF, 3.0, False, 'A roll of 8 pieces.' + EST),
    ('g_tem', 'Ounces of sushi rice per hand roll (oz)', OZF, 2.0, False, 'One cone.' + EST),
    ('g_pok', 'Ounces of sushi rice per poke bowl (oz)', OZF, 6.0, False, 'One bowl.' + EST),
    ('nori', 'Nori ($ per sheet, before tax)', USD3, 0.31, False, '100 full sheets $31.49, webstaurantstore.com, Oct 2026.'),
    ('n_nig', 'Nori sheets per nigiri', '#,##0.00', 0.05, False, 'The strip that holds the nigiri.' + EST),
    ('n_hos', 'Nori sheets per maki', '#,##0.00', 0.5, False, 'Half a sheet per roll.' + EST),
    ('n_ura', 'Nori sheets per inside-out roll', '#,##0.00', 0.5, False, 'Half a sheet per roll.' + EST),
    ('n_tem', 'Nori sheets per hand roll', '#,##0.00', 0.5, False, 'Half a sheet per cone.' + EST),
    ('n_pok', 'Nori sheets per poke bowl', '#,##0.00', 0.1, False, 'A little cut nori on top.' + EST),
    ('mt', 'Sushi rice thrown out at close (%)', PCT, 0.05, False, 'What is left and cannot be kept.' + EST),
    ('cd', 'Covers per day', ENT, 115, False, 'The case of the business plan. Replace with yours.' + EST),
    ('gc', 'Ounces of sushi rice per cover (oz)', OZF, 5.0, False, 'Pieces per cover × ounces per piece.' + EST),
    ('OZ', 'Ounces per lb (constant: 16 oz per lb)', ENT, 16, True, 'Constant: 1 lb = 16 oz. Not editable.'),
    ('GAL', 'Fluid ounces per gallon (constant: 128 fl oz per gallon)', ENT, 128, True, 'Constant: 1 gallon = 128 fl oz. Not editable.'),
]
SAL_A = [
    ('kgs', 'LB OF SUSHI RICE PER LB OF RAW RICE (lb)', '#,##0.00', '{f}*(1+{z}/{OZ})',
     'Raw to cooked factor × (1 + vinegar fl oz ÷ 16).'),
    ('cs', 'COST PER LB OF SUSHI RICE ($)', USD, '({ak}+{f}*{z}/{GAL}*{zp})/{kgs}', '(Rice + vinegar for 1 lb of raw rice) ÷ lb of sushi rice.'),
    ('csm', 'Cost per lb of sushi rice with the close discard ($)', USD, '{cs}/(1-{mt})',
     'Cost per lb of sushi rice ÷ (1 − rice thrown out). It feeds the costing.'),
    ('nigkg', 'NIGIRI PER LB OF RAW RICE (pieces)', ENT, '{OZ}*{kgs}/{g_nig}', '16 × lb of sushi rice per lb of raw rice ÷ ounces per nigiri.'),
    ('c_nig', 'Cost of the sushi rice in a nigiri ($)', USD3, '{g_nig}*{csm}/{OZ}', 'Ounces × cost per lb with discard ÷ 16.'),
    ('c_hos', 'Cost of the sushi rice in a maki ($)', USD3, '{g_hos}*{csm}/{OZ}', 'Ounces × cost per lb with discard ÷ 16.'),
    ('c_ura', 'Cost of the sushi rice in an inside-out roll ($)', USD3, '{g_ura}*{csm}/{OZ}', 'Ounces × cost per lb with discard ÷ 16.'),
    ('c_tem', 'Cost of the sushi rice in a hand roll ($)', USD3, '{g_tem}*{csm}/{OZ}', 'Ounces × cost per lb with discard ÷ 16.'),
    ('c_pok', 'Cost of the sushi rice in a poke bowl ($)', USD3, '{g_pok}*{csm}/{OZ}', 'Ounces × cost per lb with discard ÷ 16.'),
    ('o_nig', 'Cost of the nori in a nigiri ($)', USD3, '{n_nig}*{nori}', 'Sheets × price per sheet.'),
    ('o_hos', 'Cost of the nori in a maki ($)', USD3, '{n_hos}*{nori}', 'Sheets × price per sheet.'),
    ('o_ura', 'Cost of the nori in an inside-out roll ($)', USD3, '{n_ura}*{nori}', 'Sheets × price per sheet.'),
    ('o_tem', 'Cost of the nori in a hand roll ($)', USD3, '{n_tem}*{nori}', 'Sheets × price per sheet.'),
    ('o_pok', 'Cost of the nori in a poke bowl ($)', USD3, '{n_pok}*{nori}', 'Sheets × price per sheet.'),
    ('shari_dia', 'Sushi rice sold per day (lb)', '#,##0.00', '{cd}*{gc}/{OZ}', 'Covers × ounces per cover ÷ 16.'),
    ('arroz_dia', 'LB OF RAW RICE PER DAY', '#,##0.00', '{shari_dia}/{kgs}', 'Sushi rice sold ÷ lb of sushi rice per lb of raw rice.'),
    ('arroz_dia_m', 'Lb of raw rice per day counting what is thrown out', '#,##0.00', '{arroz_dia}/(1-{mt})',
     'Raw rice per day ÷ (1 − rice thrown out).'),
]
M_A, M_RA, S_A = None, None, None


def _filas(ent, sal, h0=5):
    m = OrderedDict((k[0], h0 + 1 + i) for i, k in enumerate(ent))
    ra = h0 + 1 + len(ent) + 1
    s = OrderedDict((k[0], ra + 1 + i) for i, k in enumerate(sal))
    return m, ra, s


M_A, M_RA, S_A = _filas(ENT_A, SAL_A)

#: hoja Costing: datos generales
ESC_IN = [
    ('tax', 'Sales tax included in your menu prices (%)', PCT, 0.0, 'US menus usually exclude sales tax: leave 0%. Type your rate only if your '
                                                                   'prices include it. UK: VAT is 20% dine-in and cold take-out sushi is often zero-rated.'),
    ('obj', 'Food cost target (% of the price before tax)', PCT, 0.30, 'Target for the sushi menu.' + EST),
    ('tol', 'Watch band above the target (points)', PCT, 0.05, 'Up to this margin above the target the traffic light says "Watch the cost".'),
    ('cl', 'To-go packaging per order ($, before tax)', USD, 0.60, 'Tray, lid, bag, sauces and chopsticks for one take-out order.' + EST),
    ('cp', 'Marketplace packaging per order ($, before tax)', USD, 1.00, 'Delivery container plus tamper seal.' + EST),
    ('rec', 'Marketplace price markup (%)', PCT, 0.15, 'How much you raise the price on the marketplace to compensate.' + EST),
    ('com', 'Marketplace commission (%)', PCT, 0.25, 'On the price before tax; check your contract.' + EST),
]
E = OrderedDict((k[0], 5 + i) for i, k in enumerate(ESC_IN))
E_SH, E_NO = 5 + len(ESC_IN), 5 + len(ESC_IN) + 1                                # enlaces a la hoja Sushi Rice
E_CAB = E_NO + 2
#: producto: (nombre, nº de pescado 0-4, oz de pescado, oz de arroz, hojas de nori, otros ingredientes $, precio de carta $, unidades a la semana)
PRODS = [
    ('Salmon nigiri (2 pcs)', 2, 1.0, 1.5, 0.10, 0.05, 7.00, 140),
    ('Tuna nigiri (2 pcs)', 3, 1.0, 1.5, 0.10, 0.05, 7.50, 90),
    ('Salmon sashimi (5 pcs)', 2, 2.5, 0.0, 0.0, 0.08, 14.00, 50),
    ('Salmon maki (6 pcs)', 2, 1.0, 2.5, 0.5, 0.05, 6.50, 30),
    ('California roll (8 pcs)', 0, 0.0, 3.0, 0.5, 0.75, 8.50, 120),
    ('Spicy tuna roll (8 pcs)', 3, 2.0, 3.0, 0.5, 0.25, 9.50, 100),
    ('Shrimp tempura roll (8 pcs)', 0, 0.0, 3.0, 0.5, 2.20, 11.00, 80),
    ('Salmon hand roll', 2, 1.0, 2.0, 0.5, 0.30, 7.00, 40),
    ('Salmon poke bowl (uses whole-salmon trim)', 1, 4.0, 6.0, 0.1, 1.20, 16.00, 50),
    ('Sushi platter, 12 pcs', 2, 3.0, 6.0, 0.5, 0.40, 22.00, 30),
    ('Party platter, 24 pcs', 2, 6.0, 12.0, 1.0, 0.60, 42.00, 15),
    ('Gyoza (6 pcs)', 0, 0.0, 0.0, 0.0, 1.60, 8.00, 60),
    ('Edamame', 0, 0.0, 0.0, 0.0, 0.70, 5.50, 60),
    ('Mochi ice cream (2 pcs)', 0, 0.0, 0.0, 0.0, 1.40, 6.00, 40),
]
NPR = len(PRODS)
J_NIG, J_CAL, J_SPT, J_P12, J_P24 = 0, 4, 5, 9, 10
E_P0 = E_CAB + 1
E_P1 = E_P0 + NPR - 1
E_T0 = E_P1 + 2
TOT = [
    ('ing', 'Weekly sales before tax ($)', USD, 'SUMPRODUCT(%s,%s)'),
    ('cos', 'Cost of the food sold in the week ($)', USD, 'SUMPRODUCT(%s,%s)'),
    ('fcp', 'WEIGHTED FOOD COST (%)', PCT, None),
    ('semp', 'Traffic light of the weighted food cost', 'General', None),
]
T = OrderedDict((k[0], E_T0 + 1 + i) for i, k in enumerate(TOT))
E_FIN = E_T0 + len(TOT) + 2
# columnas de la tabla
CN, CPF, CNE, CSH, CNO, COT, CPVP, CUD = 'A', 'B', 'C', 'D', 'E', 'F', 'G', 'H'
CCO, CPS, CFC, CSE, CSG, CMS, CML, CMP, CRK = 'I', 'J', 'K', 'L', 'M', 'N', 'O', 'P', 'Q'
SALP = [('c_tot', CCO, USD3), ('pvp_sin', CPS, USD), ('fc', CFC, PCT), ('sem', CSE, 'General'), ('pvp_sug', CSG, USD),
        ('sala', CMS, USD), ('llevar', CML, USD), ('plat', CMP, USD), ('rank', CRK, ENT)]

EXPUESTOS = OrderedDict([
    ('COST_LB_WHOLE_SALMON', (HP, 'B%d' % S_C['cku'])),
    ('COST_LB_SALMON_FILLET', (HP, 'C%d' % S_C['cku'])),
    ('COST_LB_TUNA_SAKU', (HP, 'D%d' % S_C['cku'])),
    ('WHOLE_PAYS_OFF_BELOW', (HP, 'B%d' % S_D['X'])),
    ('LB_SUSHI_RICE_PER_LB', (HA, 'B%d' % S_A['kgs'])),
    ('NIGIRI_PER_LB_RICE', (HA, 'B%d' % S_A['nigkg'])),
    ('RICE_PER_DAY_LB', (HA, 'B%d' % S_A['arroz_dia'])),
    ('FC_SALMON_NIGIRI', (HE, '%s%d' % (CFC, E_P0 + J_NIG))),
    ('FC_CALIFORNIA_ROLL', (HE, '%s%d' % (CFC, E_P0 + J_CAL))),
    ('FC_SPICY_TUNA_ROLL', (HE, '%s%d' % (CFC, E_P0 + J_SPT))),
    ('FC_PARTY_PLATTER_24', (HE, '%s%d' % (CFC, E_P0 + J_P24))),
    ('MARGIN_DINEIN_PLATTER24', (HE, '%s%d' % (CMS, E_P0 + J_P24))),
    ('MARGIN_MARKETPLACE_PLATTER24', (HE, '%s%d' % (CMP, E_P0 + J_P24))),
    ('FC_WEIGHTED', (HE, 'B%d' % T['fcp'])),
])


def valores_defecto():
    """Todas las entradas por defecto en un solo dict plano (claves = las de las celdas verdes)."""
    v = OrderedDict()
    for i, (_n, pr, rd, rc) in enumerate(PESCADOS):
        v['pr%d' % i], v['rd%d' % i], v['rc%d' % i] = pr, rd, rc
    for ent in (ENT_PE, ENT_A):
        for k, _r, _f2, val, _c, _n in ent:
            v[k] = val
    for k, _r, _f2, val, _n in ESC_IN:
        v[k] = val
    for j, (_nom, pf, ne, sh, no, ot, pvp, ud) in enumerate(PRODS):
        v['pf%d' % j], v['ne%d' % j], v['sh%d' % j], v['no%d' % j] = pf, ne, sh, no
        v['ot%d' % j], v['pvp%d' % j], v['ud%d' % j] = ot, pvp, ud
    return v


# ---------------------------------------------------------------------------------------------------------------------
# Cálculo en Python (espejo de las fórmulas) — fuente del autotest
# ---------------------------------------------------------------------------------------------------------------------
def calcular(v):
    r = OrderedDict()
    cols = []
    for i in range(NPE):
        c = OrderedDict()
        c['tot'] = v['rd%d' % i] + v['rc%d' % i]
        c['cku'] = v['pr%d' % i] / c['tot'] / (1 - v['m'])
        c['c_nig'] = c['cku'] * v['gn'] / v['OZ']
        c['c_sas'] = c['cku'] * v['gs'] / v['OZ']
        c['c_roll'] = c['cku'] * v['gr'] / v['OZ']
        cols.append(c)
    r['Pe'] = cols
    d = OrderedDict()
    d['lab_pz'] = v['tl'] / v['MH'] * v['hh']
    d['kg_pz'] = v['pe'] * v['rd0']
    d['lab_kg'] = d['lab_pz'] / d['kg_pz']
    d['c_ent'] = cols[0]['cku'] + d['lab_kg']
    d['X'] = cols[1]['cku'] - cols[0]['cku']
    d['sem'] = SEM_ENT if d['lab_kg'] < d['X'] else SEM_LOMO
    d['rem'] = OPC[[o for o, _t in OPC].index(v['cong'])][1]
    r['D'] = d
    a = OrderedDict()
    a['kgs'] = v['f'] * (1 + v['z'] / v['OZ'])
    a['cs'] = (v['ak'] + v['f'] * v['z'] / v['GAL'] * v['zp']) / a['kgs']
    a['csm'] = a['cs'] / (1 - v['mt'])
    a['nigkg'] = v['OZ'] * a['kgs'] / v['g_nig']
    for t in ('nig', 'hos', 'ura', 'tem', 'pok'):
        a['c_' + t] = v['g_' + t] * a['csm'] / v['OZ']
    for t in ('nig', 'hos', 'ura', 'tem', 'pok'):
        a['o_' + t] = v['n_' + t] * v['nori']
    a['shari_dia'] = v['cd'] * v['gc'] / v['OZ']
    a['arroz_dia'] = a['shari_dia'] / a['kgs']
    a['arroz_dia_m'] = a['arroz_dia'] / (1 - v['mt'])
    r['A'] = a
    r['shari'], r['nori'] = a['csm'], v['nori']
    kgc = [0.0] + [c['cku'] for c in cols]
    ps = []
    for j in range(NPR):
        q = OrderedDict()
        q['c_tot'] = v['ne%d' % j] / v['OZ'] * kgc[int(v['pf%d' % j])] + v['sh%d' % j] / v['OZ'] * r['shari'] + v['no%d' % j] * r['nori'] + v['ot%d' % j]
        q['pvp_sin'] = v['pvp%d' % j] / (1 + v['tax'])
        q['fc'] = q['c_tot'] / q['pvp_sin']
        q['sem'] = FC_OK if q['fc'] <= v['obj'] else (FC_VIG if q['fc'] <= v['obj'] + v['tol'] else FC_ALTO)
        q['pvp_sug'] = q['c_tot'] / v['obj'] * (1 + v['tax'])
        q['sala'] = q['pvp_sin'] - q['c_tot']
        q['llevar'] = q['sala'] - v['cl']
        q['plat'] = q['pvp_sin'] * (1 + v['rec']) * (1 - v['com']) - q['c_tot'] - v['cp']
        ps.append(q)
    for q in ps:
        q['rank'] = 1 + sum(1 for o in ps if o['sala'] > q['sala'])
    r['Pr'] = ps
    t = OrderedDict()
    t['ing'] = sum(q['pvp_sin'] * v['ud%d' % j] for j, q in enumerate(ps))
    t['cos'] = sum(q['c_tot'] * v['ud%d' % j] for j, q in enumerate(ps))
    t['fcp'] = 0 if t['ing'] == 0 else t['cos'] / t['ing']
    t['semp'] = FC_OK if t['fcp'] <= v['obj'] else (FC_VIG if t['fcp'] <= v['obj'] + v['tol'] else FC_ALTO)
    r['T'] = t
    return r


# ---------------------------------------------------------------------------------------------------------------------
# Helpers (del Rotisserie EN)
# ---------------------------------------------------------------------------------------------------------------------
def _f(ws, coord, v, b=False, sz=11, i=False, fmt=None, wrap=False, fill=None, locked=True):
    c = ws[coord]
    c.value = v
    c.font = Font(name=FUENTE, size=sz, bold=b, italic=i)
    if fmt:
        c.number_format = fmt
    c.alignment = Alignment(wrap_text=wrap, vertical='top')
    if fill:
        c.fill = fill
    c.protection = Protection(locked=locked)
    return c


def _nota(ws, fila, texto, desde='E', hasta='F', chars=80.0):
    if not texto:
        return
    _f(ws, '%s%d' % (desde, fila), texto, wrap=True)
    if desde != hasta:
        ws.merge_cells('%s%d:%s%d' % (desde, fila, hasta, fila))
    ws.row_dimensions[fila].height = max(ws.row_dimensions[fila].height or 0, 16, 15 * math.ceil(len(texto) / chars))


def _imprimir(ws, horizontal=False):
    ws.page_setup.paperSize = M.PAPEL_EN                                          # 1 = US Letter
    ws.page_setup.orientation = 'landscape' if horizontal else 'portrait'
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.oddFooter.center.text = FOOTER
    ws.print_options.horizontalCentered = True
    ws.protection.sheet = True                                                    # sin contraseña
    ws.sheet_view.showGridLines = False


def _dv(ws, tipo, celdas, f1=None, f2=None, op=None, titulo='Invalid value', msg=''):
    dv = DataValidation(type=tipo, formula1=f1, formula2=f2, operator=op, allow_blank=False, showErrorMessage=True,
                        errorTitle=titulo, error=msg)
    ws.add_data_validation(dv)
    for c in celdas:
        dv.add(c)


def a_excel(expr, ref):
    return re.sub(r'\{(\w+)\}', lambda m: ref[m.group(1)], expr)


def _bloque(ws, titulo_in, ent, m, ra, sal, s, neg, chars=105.0):
    """Escribe el bloque de entradas (verdes/grises) y el de resultados de una hoja de dos bloques."""
    _f(ws, 'A5', titulo_in, b=True, fill=GRIS)
    _f(ws, 'B5', 'Value', b=True, fill=GRIS)
    _f(ws, 'C5', 'Notes', b=True, fill=GRIS)
    for k, rot, fmt, val, const, nota in ent:
        f = m[k]
        _f(ws, 'A%d' % f, rot)
        if const:
            _f(ws, 'B%d' % f, val, fmt=fmt, fill=GRIS)
        else:
            _f(ws, 'B%d' % f, val, fmt=fmt, fill=VERDE, locked=False)
        _nota(ws, f, nota, desde='C', hasta='I', chars=chars)
    ref = dict((k, '$B$%d' % f) for k, f in m.items())
    ref.update((k, '$B$%d' % f) for k, f in s.items())
    for col, t in zip('ABC', ('RESULTS', 'Value', 'Notes')):
        _f(ws, '%s%d' % (col, ra), t, b=True, fill=GRIS)
    for k, rot, fmt, expr, nota in sal:
        f = s[k]
        _f(ws, 'A%d' % f, rot, b=k in neg)
        _f(ws, 'B%d' % f, '=' + a_excel(expr, ref), fmt=fmt, b=k in neg, wrap=(fmt == 'General'))
        _nota(ws, f, nota, desde='C', hasta='I', chars=chars)


def _semaforo(ws, celda, ok, ko):
    m = re.match(r'([A-Z]+)(\d+)', celda)
    ws.conditional_formatting.add(celda, FormulaRule(formula=['$%s$%s="%s"' % (m.group(1), m.group(2), ok)], fill=OK))
    ws.conditional_formatting.add(celda, FormulaRule(formula=['$%s$%s="%s"' % (m.group(1), m.group(2), ko)], fill=ROJO))


# ---------------------------------------------------------------------------------------------------------------------
# Hojas
# ---------------------------------------------------------------------------------------------------------------------
def hoja_pescado(wb):
    ws = wb.active
    ws.title = HP
    for col, w in (('A', 62), ('B', 19), ('C', 19), ('D', 19), ('E', 19), ('F', 72)):
        ws.column_dimensions[col].width = w
    _f(ws, 'A1', 'FISH YIELD CALCULATOR — COST PER USABLE LB, PER SLICE, PER CUT AND PER ROLL', b=True, sz=14)
    _f(ws, 'A2', 'Type only in the GREEN cells: everything else is calculated. Amounts are before sales tax. Every default is a Kit estimate, '
                 'editable: replace it with your distributor\'s price and, above all, with what you measure on your own fish.', wrap=True)
    ws.merge_cells('A2:F2')
    ws.row_dimensions[2].height = 32
    _f(ws, 'A3', 'FIXED NOTICE. Parasite destruction (FDA Food Code 2022, 3-402.11): fish served raw or undercooked is frozen at −4 °F or below '
                 'for 168 hours (7 days), or at −31 °F until solid and held 15 hours, or at −31 °F until solid and held 24 hours at −4 °F. '
                 'Tuna species are exempt; feed-raised aquaculture fish only with a written statement from the supplier or the farm. Keep the '
                 "records 90 days (3-402.12). Your county's rules may be stricter. The cleaning minutes (25) have no source: time it.",
       wrap=True, b=True, fill=GRIS)
    ws.merge_cells('A3:F3')
    ws.row_dimensions[3].height = 64
    # A. tres pescados y uno libre
    _f(ws, 'A5', 'A. THREE FISH AND ONE FREE COLUMN', b=True, fill=GRIS)
    for i, (nom, pr, rd, rc) in enumerate(PESCADOS):
        _f(ws, '%s5' % PC[i], nom, b=True, fill=VERDE, wrap=True, locked=False)
    _f(ws, 'F5', 'Notes', b=True, fill=GRIS)
    ws.row_dimensions[5].height = 62
    for k, rot, fmt, idx in (('pr', 'Purchase price ($/lb, before tax)', USD, 1), ('rd', 'Yield: usable loin over the weight you buy', PCT, 2),
                             ('rc', 'Trim you reuse (% of the weight you buy)', PCT, 3)):
        f = P_ROW[k]
        _f(ws, 'A%d' % f, rot)
        for i, p in enumerate(PESCADOS):
            _f(ws, '%s%d' % (PC[i], f), p[idx], fmt=fmt, fill=VERDE, locked=False)
        _nota(ws, f, P_NOTAS[k], desde='F', hasta='F', chars=75.0)
    # B. datos comunes
    _f(ws, 'A10', 'B. COMMON INPUTS', b=True, fill=GRIS)
    _f(ws, 'B10', 'Value', b=True, fill=GRIS)
    _f(ws, 'C10', 'Notes', b=True, fill=GRIS)
    for k, rot, fmt, val, const, nota in ENT_PE:
        f = M_PE[k]
        _f(ws, 'A%d' % f, rot)
        if const:
            _f(ws, 'B%d' % f, val, fmt=fmt, fill=GRIS)
        else:
            _f(ws, 'B%d' % f, val, fmt=fmt, fill=VERDE, locked=False, wrap=(k == 'cong'))
        _nota(ws, f, nota, desde='C', hasta='F', chars=125.0)
    ws.merge_cells('B%d:E%d' % (M_PE['cong'], M_PE['cong']))
    ws.row_dimensions[M_PE['cong']].height = 32
    # C. coste por lb útil
    _f(ws, 'A%d' % PE_RC, 'C. COST PER USABLE LB AND PER PIECE', b=True, fill=GRIS)
    for i in range(NPE):
        _f(ws, '%s%d' % (PC[i], PE_RC), '=%s5' % PC[i], b=True, fill=GRIS, wrap=True)
    _f(ws, 'F%d' % PE_RC, 'Notes', b=True, fill=GRIS)
    ws.row_dimensions[PE_RC].height = 62
    for k, rot, fmt, expr, nota in SAL_C:
        f = S_C[k]
        _f(ws, 'A%d' % f, rot, b=k == 'cku')
        for i in range(NPE):
            c = PC[i]
            ref = dict((kk, '$B$%d' % ff) for kk, ff in M_PE.items())
            ref.update((kk, '%s$%d' % (c, ff)) for kk, ff in P_ROW.items())
            ref.update((kk, '%s$%d' % (c, ff)) for kk, ff in S_C.items())
            _f(ws, '%s%d' % (c, f), '=' + a_excel(expr, ref), fmt=fmt, b=k == 'cku')
        _nota(ws, f, nota, desde='F', hasta='F', chars=75.0)
    # D. entero o filete
    _f(ws, 'A%d' % PE_RD, 'D. WHOLE FISH OR FILLET?', b=True, fill=GRIS)
    _f(ws, 'B%d' % PE_RD, 'Value', b=True, fill=GRIS)
    _f(ws, 'C%d' % PE_RD, 'Notes', b=True, fill=GRIS)
    ref = dict((k, '$B$%d' % f) for k, f in M_PE.items())
    ref.update((k, '$B$%d' % f) for k, f in S_D.items())
    ref.update({'rd0': '$B$%d' % P_ROW['rd'], 'cku0': '$B$%d' % S_C['cku'], 'cku1': '$C$%d' % S_C['cku'],
                'rem_r': '$B$%d:$B$%d' % (PE_O0, PE_O1), 'opc_r': '$A$%d:$A$%d' % (PE_O0, PE_O1)})
    for k, rot, fmt, expr, nota in SAL_D:
        f = S_D[k]
        neg = k in ('lab_kg', 'X', 'sem')
        _f(ws, 'A%d' % f, rot, b=neg, wrap=True)
        _f(ws, 'B%d' % f, '=' + a_excel(expr, ref), fmt=fmt, b=neg, wrap=(fmt == 'General'))
        _nota(ws, f, nota, desde='C', hasta='F', chars=125.0)
    ws.merge_cells('B%d:E%d' % (S_D['rem'], S_D['rem']))
    ws.row_dimensions[S_D['rem']].height = 48
    ws.row_dimensions[S_D['X']].height = 32
    ws.row_dimensions[S_D['sem']].height = 32
    _semaforo(ws, 'B%d' % S_D['sem'], SEM_ENT, SEM_LOMO)
    # lista del desplegable
    _f(ws, 'A%d' % PE_RO, 'PARASITE DROPDOWN LIST', b=True, fill=GRIS)
    _f(ws, 'B%d' % PE_RO, 'Reminder', b=True, fill=GRIS)
    for i, (op, rem) in enumerate(OPC):
        f = PE_O0 + i
        _f(ws, 'A%d' % f, op, wrap=True)
        _f(ws, 'B%d' % f, rem, wrap=True)
        ws.merge_cells('B%d:F%d' % (f, f))
        ws.row_dimensions[f].height = 32
    # validaciones
    _dv(ws, 'decimal', ['%s%d' % (c, P_ROW['pr']) for c in PC], f1='0', op='greaterThan', msg='Type a number greater than 0.')
    _dv(ws, 'decimal', ['%s%d' % (c, P_ROW['rd']) for c in PC], f1='0.3', f2='1', op='between', titulo='Invalid yield',
        msg='Type a value between 0.3 and 1 (0.55 = 55%).')
    _dv(ws, 'decimal', ['%s%d' % (c, P_ROW['rc']) for c in PC] + ['B%d' % M_PE['m']], f1='0', f2='0.5', op='between',
        titulo='Invalid percentage', msg='Type a percentage between 0 and 0.5 (0.05 = 5%).')
    _dv(ws, 'decimal', ['B%d' % M_PE[k] for k in ('gn', 'gs', 'gr', 'pe')], f1='0', op='greaterThan', msg='Type a number greater than 0.')
    _dv(ws, 'decimal', ['B%d' % M_PE[k] for k in ('tl', 'hh')], f1='0', op='greaterThanOrEqual', msg='Type a number greater than or equal to 0.')
    _dv(ws, 'list', ['B%d' % M_PE['cong']], f1='$A$%d:$A$%d' % (PE_O0, PE_O1), titulo='Invalid option', msg='Pick an option from the list.')
    _imprimir(ws, horizontal=True)


def hoja_arroz(wb):
    ws = wb.create_sheet(HA)
    for col, w in (('A', 60), ('B', 15), ('C', 15), ('D', 16), ('E', 14), ('F', 14), ('G', 14), ('H', 12), ('I', 15)):
        ws.column_dimensions[col].width = w
    _f(ws, 'A1', 'SUSHI RICE CALCULATOR — HOW MUCH SUSHI RICE 1 LB MAKES, WHAT IT COSTS AND HOW MUCH TO BUY PER DAY', b=True, sz=14)
    _f(ws, 'A2', 'Type only in the GREEN cells: everything else is calculated. Amounts are before sales tax. Every default is a Kit estimate, '
                 'editable: weigh your rice balls and replace the ounces with yours.', wrap=True)
    ws.merge_cells('A2:I2')
    ws.row_dimensions[2].height = 32
    _f(ws, 'A3', 'FIXED NOTICE. Time-stamp every batch (TPHC, 4 hours) or, if your county approved a variance, log the pH of every batch in your '
                 'HACCP plan records.', wrap=True, b=True, fill=GRIS)
    ws.merge_cells('A3:I3')
    ws.row_dimensions[3].height = 36
    _bloque(ws, 'A. THE RICE AND THE SUSHI RICE', ENT_A, M_A, M_RA, SAL_A, S_A, ('kgs', 'cs', 'nigkg', 'arroz_dia'))
    cel = lambda k: 'B%d' % M_A[k]
    _dv(ws, 'decimal', [cel(k) for k in ('ak', 'zp', 'f', 'g_nig', 'g_hos', 'g_ura', 'g_tem', 'g_pok', 'gc')], f1='0', op='greaterThan',
        msg='Type a number greater than 0.')
    _dv(ws, 'decimal', [cel(k) for k in ('z', 'nori', 'n_nig', 'n_hos', 'n_ura', 'n_tem', 'n_pok')], f1='0', op='greaterThanOrEqual',
        msg='Type a number greater than or equal to 0.')
    _dv(ws, 'decimal', [cel('mt')], f1='0', f2='0.5', op='between', titulo='Invalid percentage', msg='Type a percentage between 0 and 0.5 (0.05 = 5%).')
    _dv(ws, 'whole', [cel('cd')], f1='1', f2='1000000', op='between', titulo='Invalid number', msg='Type a whole number of 1 or more.')
    _imprimir(ws, horizontal=True)


def hoja_escandallo(wb):
    ws = wb.create_sheet(HE)
    for col, w in ((CN, 44), (CPF, 12), (CNE, 11), (CSH, 11), (CNO, 10), (COT, 13), (CPVP, 13), (CUD, 12), (CCO, 12), (CPS, 12), (CFC, 11),
                   (CSE, 17), (CSG, 14), (CMS, 12), (CML, 12), (CMP, 13), (CRK, 10)):
        ws.column_dimensions[col].width = w
    _f(ws, 'A1', 'SUSHI MENU COSTING — 14 PRODUCTS, PER PIECE AND PER CHANNEL', b=True, sz=14)
    _f(ws, 'A2', 'Type only in the GREEN cells (names, fish, ounces, nori, menu prices and units). Amounts are before sales tax except the menu '
                 'price. Every default is a Kit estimate, editable: type yours. The cost of the fish comes from the "%s" sheet, the cost of '
                 'the sushi rice and nori from the "%s" sheet.' % (HP, HA), wrap=True)
    ws.merge_cells('A2:%s2' % CRK)
    ws.row_dimensions[2].height = 32
    for col, t in zip('ABD', ('GENERAL INPUTS', 'Value', 'Notes')):
        _f(ws, '%s4' % col, t, b=True, fill=GRIS)
    for k, rot, fmt, val, nota in ESC_IN:
        f = E[k]
        _f(ws, 'A%d' % f, rot)
        _f(ws, 'B%d' % f, val, fmt=fmt, fill=VERDE, locked=False)
        _nota(ws, f, nota, desde='D', hasta=CRK, chars=150.0)
    _f(ws, 'A%d' % E_SH, 'Cost per lb of sushi rice with discard ($, from the Sushi Rice sheet)')
    _f(ws, 'B%d' % E_SH, "='%s'!$B$%d" % (HA, S_A['csm']), fmt=USD)
    _f(ws, 'A%d' % E_NO, 'Nori ($ per sheet, from the Sushi Rice sheet)')
    _f(ws, 'B%d' % E_NO, "='%s'!$B$%d" % (HA, M_A['nori']), fmt=USD3)
    cab = [(CN, 'PRODUCT'), (CPF, 'Fish (0 to 4)'), (CNE, 'Ounces of fish'), (CSH, 'Ounces of sushi rice'), (CNO, 'Nori sheets'),
           (COT, 'Other ingredients ($)'), (CPVP, 'Menu price ($)'), (CUD, 'Units per week'), (CCO, 'COST ($)'),
           (CPS, 'Price before tax ($)'), (CFC, 'FOOD COST (%)'), (CSE, 'Traffic light'), (CSG, 'Suggested menu price ($)'),
           (CMS, 'Margin dine-in ($)'), (CML, 'Margin take-out / pickup ($)'), (CMP, 'Margin marketplace ($)'), (CRK, 'Rank by margin')]
    for col, t in cab:
        _f(ws, '%s%d' % (col, E_CAB), t, b=True, fill=GRIS, wrap=True)
    ws.row_dimensions[E_CAB].height = 48
    kg = "'%s'!$B$%d" % (HP, M_PE['OZ'])
    fila = "'%s'!$%%s$%d" % (HP, S_C['cku'])
    elige = 'CHOOSE(%s%%d+1,0,%s)' % (CPF, ','.join(fila % c for c in PC))
    for j, (nom, pf, ne, sh, no, ot, pvp, ud) in enumerate(PRODS):
        f = E_P0 + j
        _f(ws, 'A%d' % f, nom, fill=VERDE, locked=False, b=True)
        _f(ws, '%s%d' % (CPF, f), pf, fmt='0', fill=VERDE, locked=False)
        _f(ws, '%s%d' % (CNE, f), float(ne), fmt='#,##0.0;-#,##0.0;"-"', fill=VERDE, locked=False)
        _f(ws, '%s%d' % (CSH, f), float(sh), fmt='#,##0.0;-#,##0.0;"-"', fill=VERDE, locked=False)
        _f(ws, '%s%d' % (CNO, f), float(no), fmt='#,##0.00;-#,##0.00;"-"', fill=VERDE, locked=False)
        _f(ws, '%s%d' % (COT, f), float(ot), fmt=USD, fill=VERDE, locked=False)
        _f(ws, '%s%d' % (CPVP, f), float(pvp), fmt=USD, fill=VERDE, locked=False, b=True)
        _f(ws, '%s%d' % (CUD, f), float(ud), fmt=ENT, fill=VERDE, locked=False)
        ex = {
            CCO: '%s%d/%s*%s+%s%d/%s*$B$%d+%s%d*$B$%d+%s%d' % (CNE, f, kg, elige % f, CSH, f, kg, E_SH, CNO, f, E_NO, COT, f),
            CPS: '%s%d/(1+$B$%d)' % (CPVP, f, E['tax']),
            CFC: '%s%d/%s%d' % (CCO, f, CPS, f),
            CSE: 'IF(%s%d<=$B$%d,"%s",IF(%s%d<=$B$%d+$B$%d,"%s","%s"))' % (
                CFC, f, E['obj'], FC_OK, CFC, f, E['obj'], E['tol'], FC_VIG, FC_ALTO),
            CSG: '%s%d/$B$%d*(1+$B$%d)' % (CCO, f, E['obj'], E['tax']),
            CMS: '%s%d-%s%d' % (CPS, f, CCO, f),
            CML: '%s%d-$B$%d' % (CMS, f, E['cl']),
            CMP: '%s%d*(1+$B$%d)*(1-$B$%d)-%s%d-$B$%d' % (CPS, f, E['rec'], E['com'], CCO, f, E['cp']),
            CRK: '1+COUNTIF($%s$%d:$%s$%d,">"&%s%d)' % (CMS, E_P0, CMS, E_P1, CMS, f),
        }
        for _k, col, fmt in SALP:
            neg = col in (CCO, CFC)
            _f(ws, '%s%d' % (col, f), '=' + ex[col], fmt=fmt, b=neg or col == CRK, wrap=(col == CSE))
        ws.row_dimensions[f].height = 32
    rng = '%s%d:%s%d' % (CSE, E_P0, CSE, E_P1)
    top = '%s%d' % (CSE, E_P0)
    for txt, fill in ((FC_OK, OK), (FC_VIG, AMBAR), (FC_ALTO, ROJO)):
        ws.conditional_formatting.add(rng, FormulaRule(formula=['%s="%s"' % (top, txt)], fill=fill))
    # ponderado
    _f(ws, 'A%d' % E_T0, 'WEIGHTED FOOD COST OF THE WEEK (sum of costs × units ÷ sum of prices before tax × units)', b=True, fill=GRIS)
    _f(ws, 'B%d' % E_T0, 'Value', b=True, fill=GRIS)
    rc = '$%s$%d:$%s$%d' % (CCO, E_P0, CCO, E_P1)
    rp = '$%s$%d:$%s$%d' % (CPS, E_P0, CPS, E_P1)
    ru = '$%s$%d:$%s$%d' % (CUD, E_P0, CUD, E_P1)
    exs = {'ing': 'SUMPRODUCT(%s,%s)' % (rp, ru), 'cos': 'SUMPRODUCT(%s,%s)' % (rc, ru),
           'fcp': 'IF($B$%d=0,0,$B$%d/$B$%d)' % (T['ing'], T['cos'], T['ing']),
           'semp': 'IF($B$%d<=$B$%d,"%s",IF($B$%d<=$B$%d+$B$%d,"%s","%s"))' % (
               T['fcp'], E['obj'], FC_OK, T['fcp'], E['obj'], E['tol'], FC_VIG, FC_ALTO)}
    for k, rot, fmt, _x in TOT:
        f = T[k]
        _f(ws, 'A%d' % f, rot, b=k == 'fcp')
        _f(ws, 'B%d' % f, '=' + exs[k], fmt=fmt, b=k == 'fcp', wrap=(k == 'semp'))
    ws.row_dimensions[T['semp']].height = 32
    for txt, fill in ((FC_OK, OK), (FC_VIG, AMBAR), (FC_ALTO, ROJO)):
        ws.conditional_formatting.add('B%d' % T['semp'], FormulaRule(formula=['$B$%d="%s"' % (T['semp'], txt)], fill=fill))
    _f(ws, 'A%d' % E_FIN, 'Fish: 0 = none · 1 = %s · 2 = %s · 3 = %s · 4 = %s. Other ingredients: imitation crab $1.80/lb, avocado, cucumber, sauces, '
                          'sesame, tempura shrimp ($16.44/lb, 16/20), wasabi and ginger, at Kit estimate prices. The food cost in the table is '
                          'the plate served dine-in, without drinks. On a take-out order the packaging is deducted; on the marketplace the '
                          'price is raised by the markup and the commission is deducted on that price, plus the packaging. Margins are ranked '
                          'by the dine-in margin. FIXED NOTICE: tempura and gyoza travel badly: keep them off the marketplace menu or warn '
                          'the guest. This cost is THEORETICAL: the 32%% food cost of the plan (sheet "0. Assumptions", cell B13, in the Sushi '
                          'Restaurant Financial Projections) is the real usage of the whole menu, with case shrink, trim that is not reused, '
                          'rice and nori thrown out, staff meals and comps, so this weighted theoretical cost comes out lower. Do not copy '
                          'it into the plan.' % tuple(p[0] for p in PESCADOS), i=True, wrap=True)
    ws.merge_cells('A%d:%s%d' % (E_FIN, CRK, E_FIN))
    ws.row_dimensions[E_FIN].height = 96
    ws.freeze_panes = 'B%d' % (E_CAB + 1)
    # validaciones
    _dv(ws, 'decimal', ['B%d' % E[k] for k in ('tax', 'tol', 'rec', 'com')], f1='0', f2='0.9', op='between',
        titulo='Invalid percentage', msg='Type a percentage between 0 and 0.9 (0.1 = 10%).')
    _dv(ws, 'decimal', ['B%d' % E['obj']], f1='0.01', f2='0.9', op='between', titulo='Invalid target',
        msg='The target cannot be 0 (it is a divisor): type a value between 0.01 and 0.9.')
    _dv(ws, 'decimal', ['B%d' % E['cl'], 'B%d' % E['cp']]
        + ['%s%d' % (c, f) for c in (CNE, CSH, CNO, COT, CPVP, CUD) for f in range(E_P0, E_P1 + 1)],
        f1='0', op='greaterThanOrEqual', titulo='Invalid amount', msg='Type a number greater than or equal to 0.')
    _dv(ws, 'whole', ['%s%d' % (CPF, f) for f in range(E_P0, E_P1 + 1)], f1='0', f2=str(NPE), op='between', titulo='Invalid fish',
        msg='Type 0 (none), 1, 2, 3 or 4, the number of the column on the Fish Yield sheet.')
    _imprimir(ws, horizontal=True)


def instrucciones():
    """Texto de la hoja Instructions; las cifras salen de calcular() con los valores por defecto (no se desfasan)."""
    r0 = calcular(valores_defecto())
    ck = [c['cku'] for c in r0['Pe'][:3]]
    return [
        ('INSTRUCTIONS — Sushi Fish, Rice & Costing Calculator', True),
        ('Three tools in one workbook: the "Fish Yield" sheet tells you what a usable lb of whole salmon, salmon fillet and tuna saku costs, what each '
         'slice, cut and roll is worth and whether buying the whole fish pays off; the "Sushi Rice" sheet calculates the sushi rice that 1 lb of raw '
         'rice makes, its cost, how many nigiri you get per lb and how much rice to buy per day; the "Costing" sheet calculates the cost, food cost, '
         'margin and menu price of 14 products, with the margin dine-in, take-out and on the marketplace, and the weighted food cost of your week. '
         'You type only in the GREEN cells; everything else is calculated.', False),
        ('Every default is a Kit estimate, editable: coherent examples for the case of the business plan, not anyone\'s prices. Replace them with '
         "your distributor's prices and with what you measure at your counter before you decide menu prices.", False),
        ('HOW THIS FILE FITS THE KIT', True),
        ('This calculator is one of four files: Sushi Restaurant Business Plan, Sushi Restaurant Financial Projections, Sushi Restaurant Startup '
         'Checklist and this one. The cost per piece you get here is NOT copied into the plan: the plan uses the real usage of the whole menu (see '
         '"From theoretical food cost to the plan\'s 32%" below).', False),
        ('SHEET "FISH YIELD"', True),
        ('1. Measure your yield. Weigh 3 whole salmon, break them down, weigh the usable loin from each and divide by the weight of the whole fish; '
         'type the average in "Yield". Weigh also the trim you reuse (spicy mix, maki, poke) and type it in "Trim you reuse". Weigh 10 nigiri '
         'slices and 10 sashimi cuts and use the average in ounces.', False),
        ('2. The cost per usable lb is the price ÷ (yield + trim) ÷ (1 − case shrink). It is the number that feeds the costing: with the Kit '
         'estimates it comes out at $%.2f per lb for the whole salmon, $%.2f for the fillet and $%.2f for the tuna saku.' % tuple(ck), False),
        ('3. Cleaning costs money too. The kit assumes 25 minutes per whole fish, a figure with no source: time it. The head sushi chef cost is '
         'estimated at $33.00 an hour: $5,199 a month × 12 + 10% payroll taxes ÷ 2,080 h ≈ $33/h. The sheet tells you up to what cleaning cost per '
         'usable lb buying the whole fish still pays off.', False),
        ('4. Parasite destruction (FDA Food Code 2022, 3-402.11). Fish served raw or undercooked is frozen at −4 °F or below for 168 hours (7 days), '
         'or at −31 °F until solid and held 15 hours, or at −31 °F until solid and held 24 hours at −4 °F. Tuna species and feed-raised aquaculture '
         'fish such as farmed salmon are exempt; the farmed fish only with a written statement from the supplier or the farm, and the kit also files the tuna invoice that names the species. If the fish arrives frozen, keep the '
         "supplier's written agreement or statement for every lot, showing the freezing time and temperature. Keep the records 90 days (3-402.12). "
         '"Sushi-grade" has no legal definition: a supplier label is no substitute for the freezing records or the written statement. The dropdown '
         'on the sheet is a reminder only: it changes no figure.', False),
        ('5. Scombrotoxin. Receive tuna, yellowtail and mackerel at 40 °F or below and keep them cold. Raw or undercooked fish needs the consumer '
         'advisory (3-603.11) on the menu and on your marketplace listings.', False),
        ('SHEET "SUSHI RICE"', True),
        ('1. Weigh 10 rice balls (the rice already seasoned with sushi-zu) and use the average in ounces per nigiri; do the same with a maki, an '
         'inside-out roll, a hand roll and a poke bowl.', False),
        ('2. The key results are the lb of sushi rice per lb of raw rice (%.2f with the Kit estimates), the cost per lb of sushi rice and the nigiri '
         'you get from each lb of raw rice. With them you know how much rice to buy per day for your covers.' % r0['A']['kgs'], False),
        ('3. Rice safety. Cooked sushi rice is a TCS food. Either hold it with Time as a Public Health Control (3-501.19): up to 4 hours with the '
         'time marked and a written procedure, then discard; or, if your county approved a variance, run it as acidified rice under a HACCP plan '
         '(3-502.11) and log the pH of every batch. The Food Code annex gives pH below 4.2 as an example critical limit and each county sets its '
         'own: the kit does not set a legal pH.', False),
        ('SHEET "COSTING"', True),
        ('1. Type the sales tax included in your menu prices (leave 0% if your menu is before tax), the food cost target, the watch band, the '
         'to-go and marketplace packaging, the marketplace markup and the commission.', False),
        ('2. In the table each row is a product. "Fish" is the number of the Fish Yield column it uses (1 whole salmon, 2 salmon fillet, 3 tuna '
         'saku, 4 free; 0 if it has no fish). Type the ounces of fish, the ounces of sushi rice, the nori sheets, what the other ingredients cost, '
         'the menu price and the units you sell per week.', False),
        ('3. For each product you see the cost, the price before tax, the food cost with a traffic light, the menu price you would need to reach '
         'your target and the margin dine-in, take-out / pickup and on the marketplace, with the rank of each product by dine-in margin. Below, '
         'the weighted food cost of your week. Tempura and gyoza travel badly: keep them off the marketplace menu or warn the guest.', False),
        ("FROM THEORETICAL FOOD COST TO THE PLAN'S 32%", True),
        ("The food cost on this sheet is theoretical: what you would get if everything you buy were sold exactly as in the costing. With the Kit "
         'estimates the weighted week comes out around %.0f%%. The plan works with 32%% (sheet "0. Assumptions", cell B13, in the Sushi Restaurant '
         'Financial Projections), the prudent side: between the recipe and the register you lose case shrink, trim that is not reused, ounces '
         'that weigh more than in the costing, sushi rice, nori and sauces thrown out at close, staff meals, comps and order mistakes. '
         'Your real comparison is this: the month\'s food purchases ÷ the month\'s food sales before tax. If after the first quarter it comes out '
         "above or below 32%%, update the food cost assumption in the Financial Projections." % (r0['T']['fcp'] * 100), False),
        ('PRICES, SALES TAX AND SOURCES', True),
        ('US menus usually exclude sales tax, so the calculator works before tax and the sales tax input is 0%. Type your rate only if your menu '
         'prices include it. The prices and yields come from the kit\'s market research (October 2026): rice, nori and vinegar from a restaurant '
         'supply store, salmon and tuna from secondary price reports and retail. None of the figures in this calculator is a published figure for '
         'your business: replace them with yours.', False),
    ]


def hoja_instrucciones(wb):
    ws = wb.create_sheet(HI)
    ws.column_dimensions['A'].width = 110
    f = 1
    for t, b in instrucciones():
        _f(ws, 'A%d' % f, t, b=b, sz=13 if f == 1 else 11, wrap=True)
        if not b:
            ws.row_dimensions[f].height = max(16, 15 * math.ceil(len(t) / 120.0))
        f += 2 if b and f > 1 else 1
    f += 1
    for t in ('To edit a cell that is not green: Review → Unprotect Sheet (there is no password).',
              'More templates and kits in the catalog at ' + M.URL_TIENDA, M.FIRMA_EN, VERSION_LINEA, M.AVISO_EN):
        _f(ws, 'A%d' % f, t, wrap=True, i=t == M.AVISO_EN)
        f += 1
    _imprimir(ws)


def docprops(wb):
    pr = wb.properties
    pr.title = '%s · %s' % (TITULO, PRODUCTO)
    pr.subject = '%s · v%s' % (PRODUCTO, VERSION)
    pr.keywords = KEYWORDS
    pr.description = URL
    pr.category = 'AI Chef Pro · Digital products'
    pr.creator = pr.lastModifiedBy = 'AI Chef Pro'


def generar(destino):
    wb = openpyxl.Workbook()
    hoja_pescado(wb)
    hoja_arroz(wb)
    hoja_escandallo(wb)
    hoja_instrucciones(wb)
    docprops(wb)
    os.makedirs(os.path.dirname(destino), exist_ok=True)
    wb.save(destino)
    import inject_cache as IC                                                      # noqa: E402
    nf, inj, fail = IC.inject(destino)
    print('%s: formulas=%d cache_inyectado=%d fallos_pycel=%d' % (os.path.basename(destino), nf, inj, fail))
    return fail


def valores_expuestos(ruta):
    """Lee de la CACHÉ del xlsx los valores expuestos: {nombre: valor}."""
    wv = openpyxl.load_workbook(ruta, data_only=True)
    return OrderedDict((n, wv[h][c].value) for n, (h, c) in EXPUESTOS.items())


# ---------------------------------------------------------------------------------------------------------------------
# Autotest
# ---------------------------------------------------------------------------------------------------------------------
#: restos del Taco EN / Rotisserie EN / Pizza EN, euros, español, promesas y reglamentación ES (se suma datos_sus.RX_SUS_NO)
PROHIBIDOS = [r'pizz', r'neapolit', r'\bdough\b', r'mozzarella', r'\btacos?\b', r'taquer', r'trompo', r'al pastor', r'tortilla', r'masa harina',
              r'rotisserie', r'€', r'\bIVA\b', r'APPCC', r'\bkg\b|\bgramos?\b|\bkW\b|\bkWh\b|\bml\b', r'°C',
              r'gluten[- ]free', r'authentic', r'sushiman', r'combinado', r'vitrina', r'surimi', r'RD 1021', r'RD 1086', r'anisakis',
              r'−\s?20 ?°|-\s?20 ?°', r'\bcocktails?\b|\bhappy hour\b|\btapas\b',
              r'sushi[- ]grade(?!.{0,60}no legal definition)|sashimi[- ]grade',
              r'\bfreelance', r'carnet', r'tarifa plana', r'convenio|\bSMI\b|aut[oó]nomo|Veri\*?factu']
CELDA = {}                                                                         # clave de entrada → (hoja, celda)
for _i in range(NPE):
    for _k in ('pr', 'rd', 'rc'):
        CELDA['%s%d' % (_k, _i)] = (HP, '%s%d' % (PC[_i], P_ROW[_k]))
for _k, _f2 in M_PE.items():
    CELDA[_k] = (HP, 'B%d' % _f2)
for _k, _f2 in M_A.items():
    CELDA[_k] = (HA, 'B%d' % _f2)
for _k, _f2 in E.items():
    CELDA[_k] = (HE, 'B%d' % _f2)
for _j in range(NPR):
    for _k, _c in (('pf', CPF), ('ne', CNE), ('sh', CSH), ('no', CNO), ('ot', COT), ('pvp', CPVP), ('ud', CUD)):
        CELDA['%s%d' % (_k, _j)] = (HE, '%s%d' % (_c, E_P0 + _j))


def celdas_salida(r):
    """(hoja, celda, valor esperado, etiqueta) de todas las salidas."""
    out = []
    for i, c in enumerate(r['Pe']):
        for k in S_C:
            out.append((HP, '%s%d' % (PC[i], S_C[k]), c[k], 'Pe%d.%s' % (i, k)))
    for k in S_D:
        out.append((HP, 'B%d' % S_D[k], r['D'][k], 'D.' + k))
    for k in S_A:
        out.append((HA, 'B%d' % S_A[k], r['A'][k], 'A.' + k))
    out.append((HE, 'B%d' % E_SH, r['shari'], 'shari'))
    out.append((HE, 'B%d' % E_NO, r['nori'], 'nori'))
    for j, q in enumerate(r['Pr']):
        for k, col, _fm in SALP:
            out.append((HE, '%s%d' % (col, E_P0 + j), q[k], '%s[%d]' % (k, j)))
    for k in T:
        out.append((HE, 'B%d' % T[k], r['T'][k], 'T.' + k))
    return out


def _distinto(x, v):
    if isinstance(v, str):
        return x != v
    return not isinstance(x, (int, float)) or abs(float(x) - v) > 0.01


def _checks_logica(nombre, jj, err):
    r = calcular(jj)
    for i, c in enumerate(r['Pe']):
        if abs(c['cku'] * c['tot'] * (1 - jj['m']) - jj['pr%d' % i]) > 1e-6:
            err.append('%s: cost per usable lb × (yield + trim) × (1 − shrink) ≠ price in column %d' % (nombre, i))
        if abs(c['c_nig'] * jj['OZ'] / jj['gn'] - c['cku']) > 1e-6:
            err.append('%s: slice cost ≠ cost per usable lb × ounces ÷ 16 in column %d' % (nombre, i))
    a = r['A']
    if abs(a['nigkg'] - 16 * jj['f'] * (1 + jj['z'] / 16) / jj['g_nig']) > 1e-6:
        err.append('%s: nigiri per lb ≠ 16 × f × (1 + z ÷ 16) ÷ g' % nombre)
    if abs(a['cs'] * a['kgs'] - (jj['ak'] + jj['f'] * jj['z'] / jj['GAL'] * jj['zp'])) > 1e-6:
        err.append('%s: cost per lb of sushi rice × lb of sushi rice ≠ cost of the rice and the vinegar' % nombre)
    if abs(a['arroz_dia'] * a['kgs'] * 16 - jj['cd'] * jj['gc']) > 1e-6:
        err.append('%s: rice per day × sushi rice per lb ≠ covers × ounces' % nombre)
    for j in (J_P24, J_P12):
        q = r['Pr'][j]
        if not q['plat'] < q['llevar'] < q['sala']:
            err.append('%s: marketplace < take-out < dine-in margin fails on platter %d' % (nombre, j))
    for j, q in enumerate(r['Pr']):
        if not q['plat'] < q['llevar'] < q['sala']:
            err.append('%s: marketplace < take-out < dine-in margin fails on product %d' % (nombre, j))
    if abs(r['D']['X'] - (r['Pe'][1]['cku'] - r['Pe'][0]['cku'])) > 1e-9:
        err.append('%s: X ≠ cost per usable lb of the fillet − of the whole fish' % nombre)


def autotest(destino):
    from pycel import ExcelCompiler
    err, n = [], 0
    wv = openpyxl.load_workbook(destino, data_only=True)
    wf = openpyxl.load_workbook(destino)
    base = valores_defecto()
    j2 = dict(base, pr0=8.5, rd0=0.5, rc0=0.1, pr1=19.0, rc1=0.0, pr2=28.0, rd2=0.85, pr3=15.0, rd3=0.7, m=0.08, gn=0.6, gs=0.7, gr=2.0,
              tl=35.0, hh=28.0, pe=9.0, cong=OPC[1][0], ak=1.30, f=2.0, z=1.2, zp=24.0, g_nig=0.8, g_ura=3.5, nori=0.40, n_ura=0.6, mt=0.1,
              cd=100, gc=6.0, tax=0.08, obj=0.28, tol=0.04, cl=0.8, cp=1.2, rec=0.1, com=0.3, pf3=2, ne3=2.0, sh6=3.5, ot9=0.5,
              pvp10=48.0, ud0=100, ud10=40)
    j3 = dict(base, pr0=6.5, rd0=0.6, pr1=22.0, rd1=0.9, pr2=20.0, pr3=40.0, rd3=0.6, rc3=0.1, m=0.02, gn=0.4, tl=15.0, hh=24.0, pe=12.0,
              cong=OPC[2][0], ak=0.9, f=2.5, z=2.0, zp=18.0, g_nig=0.7, nori=0.25, mt=0.0, cd=60, gc=4.0, obj=0.35, tol=0.1, cl=0.4, cp=0.8,
              rec=0.2, com=0.35, pf0=4, ne0=1.0, pf10=2, ne10=8.0, sh10=14.0, pvp9=26.0, pvp10=40.0, ud5=0, ud12=200)
    # 1. caché y pycel del juego por defecto; pycel de los otros dos con set_value
    xl = ExcelCompiler(filename=destino)
    r0 = calcular(base)
    for hoja, c, v, lab in celdas_salida(r0):
        n += 1
        for et, x in (('pycel', xl.evaluate("'%s'!%s" % (hoja, c))), ('cache', wv[hoja][c].value)):
            if _distinto(x, v):
                err.append('defaults · %s!%s (%s): %s %r ≠ expected %r' % (hoja, c, lab, et, x, v))
    for nombre, jj in (('set 2', j2), ('set 3', j3)):
        r = calcular(jj)
        xl2 = ExcelCompiler(filename=destino)
        for hoja, c, v, lab in celdas_salida(r0):
            xl2.evaluate("'%s'!%s" % (hoja, c))
        for k, v in jj.items():
            if v != base[k]:
                h, c = CELDA[k]
                xl2.set_value("'%s'!%s" % (h, c), v)
        for hoja, c, v, lab in celdas_salida(r):
            n += 1
            x = xl2.evaluate("'%s'!%s" % (hoja, c))
            if _distinto(x, v):
                err.append('%s · %s!%s (%s): pycel %r ≠ expected %r' % (nombre, hoja, c, lab, x, v))
    # 2. identidades de la lógica en los tres juegos
    for nombre, jj in (('defaults', base), ('set 2', j2), ('set 3', j3)):
        n += 12
        _checks_logica(nombre, jj, err)
    # identidades sobre la hoja viva (pycel, valores por defecto)
    ev = lambda h, c: xl.evaluate("'%s'!%s" % (h, c))
    for i in range(NPE):
        n += 1
        pr_, rd_, rc_ = (ev(HP, '%s%d' % (PC[i], P_ROW[k])) for k in ('pr', 'rd', 'rc'))
        m_ = ev(HP, 'B%d' % M_PE['m'])
        if abs(ev(HP, '%s%d' % (PC[i], S_C['cku'])) - pr_ / (rd_ + rc_) / (1 - m_)) > 1e-6:
            err.append('Fish Yield sheet, column %s: cost per usable lb ≠ price ÷ (y + t) ÷ (1 − s)' % PC[i])
    kgs, ng, gn = ev(HA, 'B%d' % S_A['kgs']), ev(HA, 'B%d' % S_A['nigkg']), ev(HA, 'B%d' % M_A['g_nig'])
    fz, zz = ev(HA, 'B%d' % M_A['f']), ev(HA, 'B%d' % M_A['z'])
    n += 1
    if abs(ng - 16 * fz * (1 + zz / 16) / gn) > 1e-6 or abs(kgs * 16 / gn - ng) > 1e-6:
        err.append('Sushi Rice sheet: nigiri per lb ≠ 16 × f × (1 + z ÷ 16) ÷ g')
    # 3. valores con sentido
    nig, spt = r0['Pr'][J_NIG], r0['Pr'][J_SPT]
    if not 0.18 <= nig['fc'] <= 0.30:
        err.append('salmon nigiri food cost %.1f %% outside 18-30 %%' % (nig['fc'] * 100))
    if not spt['fc'] > base['obj']:
        err.append('spicy tuna roll food cost %.1f %% is not above the target' % (spt['fc'] * 100))
    if not 0.22 <= r0['T']['fcp'] <= 0.32:
        err.append('default weighted food cost %.1f %% outside 22-32 %%' % (r0['T']['fcp'] * 100))
    ck = [c['cku'] for c in r0['Pe'][:3]]
    if any(abs(a_ - b_) > 0.01 for a_, b_ in zip(ck, (11.28, 23.92, 25.48))):
        err.append('default cost per usable lb %r ≠ 11.28 · 23.92 · 25.48' % [round(x, 3) for x in ck])
    A_ = r0['A']
    if abs(A_['kgs'] - 2.41) > 0.005 or abs(A_['nigkg'] - 51) > 0.5:
        err.append('default sushi rice: 2.41 lb per lb and 51 nigiri per lb (got %r)' % [round(A_[k], 3) for k in ('kgs', 'nigkg')])
    if r0['D']['sem'] != SEM_ENT:
        err.append('whole fish or fillet at defaults: "%s" (expected the whole fish pays off)' % r0['D']['sem'])
    if calcular(dict(base, tl=2400.0))['D']['sem'] != SEM_LOMO:
        err.append('the verdict does not switch to "Buy the fillet" with 2,400 cleaning minutes')
    if sorted(q['rank'] for q in r0['Pr']) != list(range(1, NPR + 1)) and len(set(q['rank'] for q in r0['Pr'])) < NPR - 2:
        err.append('the ranks by margin are not coherent')
    for j in (J_P12, J_P24):
        q = r0['Pr'][j]
        n += 1
        if not q['plat'] < q['llevar'] < q['sala']:
            err.append('platter %d: marketplace < take-out < dine-in margin fails at defaults' % j)
    # 4. valores expuestos: caché frente a calcular()
    p24 = r0['Pr'][J_P24]
    esperados = {'COST_LB_WHOLE_SALMON': r0['Pe'][0]['cku'], 'COST_LB_SALMON_FILLET': r0['Pe'][1]['cku'], 'COST_LB_TUNA_SAKU': r0['Pe'][2]['cku'],
                 'WHOLE_PAYS_OFF_BELOW': r0['D']['X'], 'LB_SUSHI_RICE_PER_LB': A_['kgs'], 'NIGIRI_PER_LB_RICE': A_['nigkg'],
                 'RICE_PER_DAY_LB': A_['arroz_dia'], 'FC_SALMON_NIGIRI': nig['fc'], 'FC_CALIFORNIA_ROLL': r0['Pr'][J_CAL]['fc'],
                 'FC_SPICY_TUNA_ROLL': spt['fc'], 'FC_PARTY_PLATTER_24': p24['fc'], 'MARGIN_DINEIN_PLATTER24': p24['sala'],
                 'MARGIN_MARKETPLACE_PLATTER24': p24['plat'], 'FC_WEIGHTED': r0['T']['fcp']}
    expuestos = valores_expuestos(destino)
    for nombre, e_ in esperados.items():
        n += 1
        if _distinto(expuestos[nombre], e_):
            err.append('exposed %s: cache %r ≠ %r' % (nombre, expuestos[nombre], e_))
    # 5. formato, protección, DV, fórmulas, textos
    if motor.gate_formatos(wf, FICHERO):
        err.append('gate_formatos: %r' % motor.gate_formatos(wf, FICHERO)[:3])
    dvh = dict((w.title, dict((str(rng), dv) for dv in w.data_validations.dataValidation for rng in dv.sqref.ranges))
               for w in (wf[HP], wf[HA], wf[HE]))
    for hoja, k, tipo in ((HP, 'pr0', 'decimal'), (HP, 'rd1', 'decimal'), (HP, 'm', 'decimal'), (HP, 'cong', 'list'), (HA, 'ak', 'decimal'),
                          (HA, 'cd', 'whole'), (HE, 'obj', 'decimal'), (HE, 'tax', 'decimal'), (HE, 'com', 'decimal'), (HE, 'pf0', 'whole'),
                          (HE, 'ud0', 'decimal')):
        c = CELDA[k][1]
        dv = dvh[hoja].get(c)
        if dv is None or dv.type != tipo:
            err.append('DV missing or of another type in %s!%s (expected %s)' % (hoja, c, tipo))
    textos = []
    for w in wf.worksheets:
        if w.page_setup.paperSize != 1 or not w.sheet_properties.pageSetUpPr.fitToPage or not w.protection.sheet \
                or w.oddFooter.center.text != FOOTER:
            err.append('%s: Letter + fit + protection + footer' % w.title)
        if w.protection.password:
            err.append('%s: protected with a password' % w.title)
        for row in w.iter_rows():
            for c in row:
                v = c.value
                if v == '':
                    err.append('%s!%s: empty string' % (w.title, c.coordinate))
                if isinstance(v, str) and c.data_type != 'f':
                    textos.append((w.title, c.coordinate, v))
                if isinstance(v, str) and c.data_type == 'f':
                    textos += [(w.title, c.coordinate + ' (formula)', lit) for lit in re.findall(r'"((?:[^"]|"")*)"', v)
                               if any(ch.isalpha() for ch in lit)]
                if c.data_type == 'f':
                    if wv[w.title][c.coordinate].value is None:
                        err.append('%s!%s: formula without cache' % (w.title, c.coordinate))
                    sin_refs = re.sub(r'"[^"]*"', '', re.sub(r"'[^']+'!|\$?[A-Z]{1,2}\$?\d+", '', str(c.value)))
                    for num in re.findall(r'(?<![A-Za-z])\d+\.?\d*', sin_refs):
                        if num not in ('0', '1'):
                            err.append('%s!%s: typed number %s in %s' % (w.title, c.coordinate, num, c.value))
                if c.fill.fgColor.rgb in ('00E8F5E9', 'FFE8F5E9') and c.protection.locked:
                    err.append('%s!%s: locked green cell' % (w.title, c.coordinate))
                if c.fill.fgColor.rgb in ('00E8F5E9', 'FFE8F5E9') and c.data_type == 'f':
                    err.append('%s!%s: green cell with a formula' % (w.title, c.coordinate))
                fmt = c.number_format or ''
                if '€' in fmt or '$' in fmt:
                    err.append('%s!%s: currency symbol in the format (D8)' % (w.title, c.coordinate))
        for dv_ in w.data_validations.dataValidation:
            for t in (dv_.error, dv_.errorTitle, dv_.prompt, dv_.promptTitle):
                if t:
                    textos.append((w.title, 'DV %s' % dv_.sqref, t))
    for hoja, ks in ((HP, ('OZ', 'MH')), (HA, ('OZ', 'GAL'))):
        for k in ks:
            c = wf[hoja]['B%d' % (M_PE if hoja == HP else M_A)[k]]
            if c.fill.fgColor.rgb not in ('00F2F2F2', 'FFF2F2F2') or c.protection.locked is not True:
                err.append('gray constant %s badly marked or unlocked' % k)
    p = wf.properties
    for campo in ('title', 'subject', 'keywords', 'description', 'category'):
        textos.append(('docProps', campo, getattr(p, campo) or ''))
    rx_sus = []
    try:
        import datos_sus as D
        rx_sus = [(x[0], x[2] if len(x) > 2 else None) for x in D.RX_SUS_NO]
    except Exception as e:                                                         # noqa: BLE001
        print('WARNING: datos_sus.RX_SUS_NO could not be imported (%s)' % e)
    n_es = n_pz = 0
    for h, c, t in textos:
        r = [x for x in M.restos_espanol(t) if not x[0].startswith('oficio')]
        if r:
            n_es += 1
            err.append('%s!%s: Spanish leftovers %s' % (h, c, r[:1]))
        if M.no_latinos(t):
            err.append('%s!%s: non-Latin characters' % (h, c))
        for rx in PROHIBIDOS:
            if re.search(rx, t, re.I if not rx.startswith('(?i)') else 0):
                n_pz += 1
                err.append('%s!%s: forbidden /%s/ in %r' % (h, c, rx, t[:60]))
        for rx, exc in rx_sus:
            if re.search(rx, t) and not (exc and re.search(exc, t)):
                n_pz += 1
                err.append('%s!%s: RX_SUS_NO /%s/ in %r' % (h, c, rx, t[:60]))
        if re.search(r'(?i)\bvalue\b.*\$\d|\bsave \$|total value|\bwas \$|\d+\s?% off\b', t):
            err.append('%s!%s: price anchor' % (h, c))
    todo = '\n'.join(t for _h, _c, t in textos)
    for frase in ('16 oz per lb', '128 fl oz per gallon', '60 minutes per hour', 'Sales tax included in your menu prices (%)', '−4 °F', '−31 °F',
                  '3-402.11', '3-402.12', '3-501.19', '3-502.11', '3-603.11', "Your county's rules", SEM_ENT, SEM_LOMO, FC_OK, FC_VIG, FC_ALTO,
                  'Kit estimate, editable', "Kit estimate (Oct 2026): replace with your distributor's price", 'No source: time it',
                  'imitation crab', 'Sushi Restaurant Financial Projections', 'Sushi Restaurant Startup Checklist',
                  'Sushi Restaurant Business Plan', 'Tempura and gyoza travel badly', 'tempura and gyoza travel badly',
                  'Time-stamp every batch (TPHC, 4 hours)', 'has no legal definition'):
        if frase not in todo:
            err.append('missing text %r' % frase)
    if M.FIRMA_EN not in todo or VERSION_LINEA not in todo or M.AVISO_EN not in todo:
        err.append('missing signature, version line or notice')
    if wf[HE]['A5'].value != 'Sales tax included in your menu prices (%)' or wf[HE]['B5'].value != 0:
        err.append('Costing!B5 is not the sales tax = 0 %')
    if wf.sheetnames != [HP, HA, HE, HI]:
        err.append('sheet names %r' % wf.sheetnames)
    if NPR != 14:
        err.append('not 14 products')
    if p.creator != 'AI Chef Pro' or p.description != URL or 'AI Chef Pro' not in (p.keywords or '') \
            or not (p.title or '').startswith(TITULO):
        err.append('docProps: creator=%r title=%r description=%r keywords=%r' % (p.creator, p.title, p.description, p.keywords))
    for x in err[:80]:
        print('ERROR', x)
    print('autotest sushi fish, rice & costing calculator: %d checks (pycel + cache at defaults; pycel in 2 more sets; identities) '
          'over %d texts · %s' % (n, len(textos), 'GREEN' if not err else 'RED (%d errors)' % len(err)))
    print('  fish ($/usable lb): ' + ' · '.join('%s %.2f' % (PESCADOS[i][0], c['cku']) for i, c in enumerate(r0['Pe'])))
    d = r0['D']
    print('  cleaning: $%.3f per usable lb · the whole fish pays off up to $%.3f per usable lb · %s' % (d['lab_kg'], d['X'], d['sem']))
    print('  rice: %.3f lb sushi rice per lb raw · $%.3f per lb ($%.3f with discard) · %.1f nigiri per lb raw · %.2f lb raw per day'
          % (A_['kgs'], A_['cs'], A_['csm'], A_['nigkg'], A_['arroz_dia']))
    for nombre, (h, c) in EXPUESTOS.items():
        print('  %-30s %s!%-5s %s' % (nombre, h, c, expuestos[nombre]))
    for j, p_ in enumerate(PRODS):
        q = r0['Pr'][j]
        print('  %-44s cost $%.3f · food cost %5.1f %% · %-14s · margin dine-in %.2f / take-out %.2f / marketplace %.2f · rank %d · %d units'
              % (p_[0], q['c_tot'], q['fc'] * 100, q['sem'], q['sala'], q['llevar'], q['plat'], q['rank'], p_[7]))
    print('  weighted: sales $%.2f · cost $%.2f · food cost %.1f %% · %s' % (r0['T']['ing'], r0['T']['cos'], r0['T']['fcp'] * 100, r0['T']['semp']))
    return 1 if err else 0


if __name__ == '__main__':
    ap = argparse.ArgumentParser()
    ap.add_argument('--salida', required=True, help='output folder')
    ap.add_argument('--autotest', action='store_true')
    a = ap.parse_args()
    destino = os.path.join(a.salida, FICHERO)
    if not a.autotest and generar(destino):
        sys.exit('ERROR: fallos de pycel en inject_cache')
    sys.exit(autotest(destino))
