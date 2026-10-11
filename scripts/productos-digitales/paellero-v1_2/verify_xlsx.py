import openpyxl, glob, io, subprocess, zipfile, logging
logging.disable(logging.CRITICAL)
from pycel import ExcelCompiler
R='/root/wt-paellero'; P='astro-site/public/dl/plan-negocio-paellero-eventos/'
for f in sorted(glob.glob(f'{R}/{P}*.xlsx')):
    n=f.split('/')[-1]; assert zipfile.ZipFile(f).testzip() is None
    new=openpyxl.load_workbook(f); old=openpyxl.load_workbook(io.BytesIO(subprocess.check_output(['git','-C',R,'show',f'HEAD:{P}{n}'])))
    dat=openpyxl.load_workbook(f,data_only=True)
    ren={'96 proveedores':'88 proveedores'}
    diffs=[]
    for wo in old:
        wn=new[ren.get(wo.title,wo.title)]
        for r in wo.iter_rows():
            for c in r:
                v=wn[c.coordinate].value
                if c.value!=v: diffs.append((wo.title,c.coordinate,c.value,v))
    print(n,'celdas con valor distinto:',[(d[1],) for d in diffs if not (d[3]=='☐')][:20],'| ☐ nuevas',sum(1 for d in diffs if d[3]=='☐'))
    ec=ExcelCompiler(filename=f); errs=0;nf=0;nocache=[]
    for ws in new:
        for r in ws.iter_rows():
            for c in r:
                if isinstance(c.value,str) and c.value.startswith('='):
                    nf+=1; a=f"'{ws.title}'!{c.coordinate}"; v=ec.evaluate(a)
                    if isinstance(v,str) and v.startswith('#'): errs+=1; print('ERR',a,v)
                    if dat[ws.title][c.coordinate].value is None: nocache.append(c.coordinate)
    print('  formulas',nf,'errores',errs,'sin cache',nocache,'sheets',new.sheetnames)
    if n.startswith('calc'):
        s="'Calculadora pricing'!"
        for pax in (None,0,80):
            ec=ExcelCompiler(filename=f); ec.evaluate(s+'B30'); ec.set_value(s+'B6',pax); print('  pax',pax,[ec.evaluate(s+x) for x in('B29','B30','B31','B32')])
        print('  cache B29,B30,B31:',[dat['Calculadora pricing'][x].value for x in('B29','B30','B31')])
    if n.startswith('plan-fin'):
        d=dat; print('  cache ingresos P&L/Proy, BE:',d['P&L Año 1']['N7'].value,d['Proyección 3 años']['B11'].value,d['Break-even']['B9'].value)
    if n.startswith('check'):
        w=dat.active; print('  contador cache',w['H3'].value,w['J3'].value,w['E3'].value)
