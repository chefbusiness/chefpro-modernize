#!/usr/bin/env python3
"""fix_docx.py — arreglos mínimos a los 7 docx del Plan de Negocio Paellero (v1.2).
Dry-run por defecto (escribe en /root/paellero-dry/); APPLY=1 escribe sobre los originales. Idempotente.
Sustituciones que conservan formato: se hacen dentro del run o fusionando solo los runs afectados."""
import os, copy, sys
from docx import Document
from docx.shared import Cm

BASE = os.path.join(os.path.dirname(os.path.abspath(__file__)), '..', '..', '..', 'astro-site', 'public', 'dl', 'plan-negocio-paellero-eventos')
APPLY = os.environ.get('APPLY') == '1'
OUT = BASE if APPLY else '/root/paellero-dry'
os.makedirs(OUT, exist_ok=True)

TITULOS = {
 'plan-de-negocio-paellero-eventos.docx': 'Plan de Negocio Paellero / Paella para Eventos',
 'catalogo-equipamiento-utensilios-menaje-paellero.docx': 'Catálogo Equipamiento, Utensilios y Menaje Paellero',
 'carta-12-paellas-arroces.docx': 'Carta 12 Paellas y Arroces',
 'modelos-contrato-paellero-eventos.docx': 'Modelos de Contrato Paellero / Paella para Eventos',
 '10-experiencias-tematicas-paellero.docx': '10 Experiencias Temáticas Paellero',
 'manual-tecnico-paellero-profesional.docx': 'Manual Técnico Paellero Profesional',
 'guia-tipos-paellero-roadmap-food-truck.docx': 'Guía 8 Sistemas de Cocción + Roadmap Food Truck',
}

def all_pars(doc):
    for p in doc.paragraphs: yield p
    for t in doc.tables:
        for r in t.rows:
            for c in r.cells:
                for p in c.paragraphs: yield p
    for s in doc.sections:
        for p in s.header.paragraphs + s.footer.paragraphs: yield p

def replace_in_par(p, old, new):
    """Sustituye old por new conservando el formato del run donde empieza old."""
    full = ''.join(r.text for r in p.runs)
    i = full.find(old)
    if i < 0: return False
    j = i + len(old); pos = 0; first = True
    for r in p.runs:
        a, b = pos, pos + len(r.text); pos = b
        if b <= i or a >= j: continue
        s, e = max(i, a) - a, min(j, b) - a
        if first:
            r.text = r.text[:s] + new + r.text[e:]; first = False
        else:
            r.text = r.text[:s] + r.text[e:]
    return True

def set_par(p, new):
    runs = p.runs
    runs[0].text = new
    for r in runs[1:]: r.text = ''

# (viejo, nuevo) por fichero; sustituciones de fragmento
SUBS = {
 'plan-de-negocio-paellero-eventos.docx': [('permite ticket medios de', 'permite tickets medios de'),
    ('(mix 30/70)', '(mix 35/65)'),
    ('Plan PIDOM.', 'Plan INFOMA.'),
    ("(Pla d'actuació per al risc d'incendis forestals)", "(Pla especial d'emergències per incendis forestals de Catalunya)"),
    ('Plan IBANAT.', 'Plan INFOBAL.'),
    ("PATRICOVA + Pla Especial Front d'Incendis Forestals.", "Pla Especial davant del risc d'incendis forestals de la Comunitat Valenciana."),
    ('Plan SOSDeiak.', 'Plan Especial de Emergencias por Riesgo de Incendios Forestales de Euskadi.'),
    ('INFOCYL.', 'INFOCAL.')],
 'manual-tecnico-paellero-profesional.docx': [
    ('(8-12 min, fuego alto)', '(12-14 min, fuego alto)'),
    ('Mantener 10-12 min a fuego alto', 'Mantener 12-14 min a fuego alto')],
 'guia-tipos-paellero-roadmap-food-truck.docx': [
    ('Perdida del componente cultural', 'Pérdida del componente cultural'),
    ('Madrid: Decreto 1/2014 + ordenanza municipal.', 'Madrid: normativa autonómica de comercio ambulante y ordenanza municipal.'),
    ('Ley 1/2011 + ordenanza.', 'Ley 3/2011 + Decreto 65/2012 + ordenanza.'),
    ('Ley 9/1988 + decreto desarrollo.', 'Decreto Legislativo 2/2012 (texto refundido de la Ley del Comercio Ambulante).'),
    ('Baleares: Decret 130/2010.', 'Baleares: Llei 11/2014 de comerç de les Illes Balears.'),
    ('Galicia: Decreto 84/2009.', 'Galicia: Lei 13/2010 de comercio interior de Galicia.'),
    ('Castilla y León: Decreto 187/2008.', 'Castilla y León: normativa autonómica de comercio ambulante y ordenanza municipal.'),
    ('La inversión está en torno a 30.000-75.000 € según escenario', 'La inversión está en torno a 26.000-86.000 € según escenario (ver tabla del apartado 2)')],
 'carta-12-paellas-arroces.docx': [('8,90 € / pax (food cost)', '8,90 € / pax'), ('26-32 € / pax (PVP)', '26-32 € / pax')],
 'modelos-contrato-paellero-eventos.docx': [
    ('adapta nombre, fecha, importe e importes', 'adapta nombre, fecha, importe y demás datos'),
    ('La factura del Paellero se emitirá a la Agencia con IVA al 21 %.',
     'La factura del Paellero se emitirá a la Agencia con IVA al 10 % (servicio de hostelería, art. 91.Uno.2.2.º de la Ley 37/1992 del IVA).'),
 ],
}
# párrafos completos del contrato B2C (cláusulas 4b y 4c)
C_OLD = [
 'Cancelación con más de 30 días de antelación: devolución del 100 % del anticipo.',
 'Cancelación entre 30 y 15 días antes del evento: retención del 50 % del anticipo.',
 'Cancelación con menos de 15 días: pérdida total del anticipo.',
 'Cancelación con menos de 7 días: el Cliente abonará el 100 % del importe acordado.',
]
C_NEW = [
 'Cancelación por el Cliente con más de 60 días de antelación: el Cliente abonará el 30 % del importe total del servicio (cubierto con el anticipo ya entregado).',
 'Cancelación por el Cliente entre 60 y 30 días antes del evento: el Cliente abonará el 50 % del importe total del servicio.',
 'Cancelación por el Cliente con menos de 30 días de antelación: el Cliente abonará el 75 % del importe total del servicio.',
 'Cancelación por el Cliente en las 72 horas previas al evento: el Cliente abonará el 100 % del importe total del servicio.',
]
C_PREST = ('Cancelación por el Paellero sin causa justificada: el Paellero devolverá al Cliente todas las cantidades recibidas y le abonará, '
           'además, una indemnización equivalente a la que correspondería al Cliente según la escala anterior en función de la antelación '
           'con la que se produzca la cancelación. Quedan exceptuados los supuestos de fuerza mayor de la cláusula quinta.')
J_OLD = 'Para cualquier controversia derivada del presente contrato, ambas partes se someten a los Juzgados y Tribunales de [Ciudad], renunciando expresamente a su propio fuero.'
J_NEW = ('Para cualquier controversia derivada del presente contrato, al tratarse de un contrato con consumidor serán competentes los Juzgados y Tribunales '
         'del domicilio del Cliente. Si el Cliente contrata como empresa, serán competentes los Juzgados y Tribunales de [Ciudad] que las partes pacten.')

def contrato(doc):
    n = 0
    for p in list(doc.paragraphs):
        t = p.text.strip()
        if t in C_OLD:
            set_par(p, C_NEW[C_OLD.index(t)]); n += 1
            if C_OLD.index(t) == 3:
                q = copy.deepcopy(p._p); p._p.addnext(q)
                from docx.text.paragraph import Paragraph
                set_par(Paragraph(q, p._parent), C_PREST)
        elif t == J_OLD:
            set_par(p, J_NEW); n += 1
    return n

NOTA_ALIOLI = ('Seguridad alimentaria: en eventos y servicios a terceros elabora el alioli con huevo pasteurizado (ovoproducto) '
               'o en su versión tradicional sin huevo (solo ajo y aceite), y mantenlo refrigerado por debajo de 8 °C hasta el servicio.')

def manual(doc):
    import re
    from docx.text.paragraph import Paragraph
    n = 0
    pars = doc.paragraphs
    for i, p in enumerate(pars):
        if p.text.startswith('3 dientes ajo + 1 huevo') and not pars[i + 1].text.startswith('Seguridad alimentaria'):
            q = copy.deepcopy(p._p); p._p.addnext(q); set_par(Paragraph(q, p._parent), NOTA_ALIOLI); n += 1
            break
    for t in doc.tables:
        hdr = [c.text.strip() for c in t.rows[0].cells]
        if 'Min total' not in hdr or 'Tiempos por fase' not in hdr: continue
        a, b = hdr.index('Min total'), hdr.index('Tiempos por fase')
        for r in t.rows[1:]:
            tot = str(sum(int(x) for x in re.findall(r'\d+', r.cells[b].text)))
            if r.cells[a].text.strip() != tot:
                set_par(r.cells[a].paragraphs[0], tot); n += 1
    return n

def main():
    for f, titulo in TITULOS.items():
        doc = Document(os.path.join(BASE, f)); cambios = []
        for old, new in SUBS.get(f, []):
            for p in all_pars(doc):
                if replace_in_par(p, old, new): cambios.append(old[:40])
        if f.startswith('modelos-contrato'):
            cambios += ['contrato'] * contrato(doc)
        if f.startswith('manual-tecnico'):
            cambios += ['manual'] * manual(doc)
        for s in doc.sections:
            s.page_width, s.page_height = Cm(21.0), Cm(29.7)
        doc.core_properties.author = 'AI Chef Pro'
        doc.core_properties.title = titulo + ' · v1.2'
        doc.save(os.path.join(OUT, f))
        print(f, len(cambios), 'cambios', cambios)

main()
