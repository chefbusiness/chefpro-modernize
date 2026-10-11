#!/usr/bin/env python3
"""fix_xlsx.py — arreglos v1.2 de los 4 Excel del Plan de Negocio Paellero.
Dry-run por defecto (escribe en /root/paellero-dry/); real con APPLY=1.
Idempotente: si no hay cambios no guarda (y no toca la cache). Tras guardar corre inject_cache.
Ejecutar en el VPS: /root/venv-guias/bin/python fix_xlsx.py"""
import os, shutil, sys, re, copy
import openpyxl
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import PatternFill, Font, Alignment

ROOT = '/root/wt-paellero'
D = f'{ROOT}/astro-site/public/dl/plan-negocio-paellero-eventos/'
APPLY = os.environ.get('APPLY') == '1'
OUT = D if APPLY else '/root/paellero-dry/'
os.makedirs(OUT, exist_ok=True)
EUR = '#,##0.00 €'
CNT = '#,##0'

class Ctx:
    def __init__(s, name):
        s.name, s.log = name, []
    def val(s, ws, coord, v):
        c = ws[coord]
        if c.value != v:
            s.log.append(f'{ws.title}!{coord}: {c.value!r} -> {v!r}'); c.value = v
    def fmt(s, ws, coord, f):
        c = ws[coord]
        if c.number_format != f:
            s.log.append(f'{ws.title}!{coord} formato {c.number_format} -> {f}'); c.number_format = f
    def dv(s, ws, sqref, formula1, titulo, msg, typ='list'):
        for d in ws.data_validations.dataValidation:
            if str(d.sqref) == sqref and d.formula1 == formula1: return
        v = DataValidation(type=typ, formula1=formula1, allow_blank=(typ != 'list' or True),
                           showErrorMessage=True, errorStyle='stop', errorTitle=titulo, error=msg,
                           showDropDown=False)  # showDropDown=False = flecha visible
        v.add(sqref); ws.add_data_validation(v)
        s.log.append(f'{ws.title}: validación {sqref} = {formula1}')

def lista(*vals): return '"' + ','.join(vals) + '"'
SINO = lista('sí', 'no')

def calculadora(wb, c):
    ws = wb['Calculadora pricing']
    c.val(ws, 'B30', '=IFERROR(B29/B6,"")')
    c.fmt(ws, 'B31', EUR)
    c.dv(ws, 'B5', lista('B2C', 'B2B'), 'Tipo de cliente', 'Elige B2C o B2B de la lista.')
    c.dv(ws, 'B8', lista('Valenciana', 'Marisco', 'Senyoret', 'Negro', 'Paella mixta', 'Banda', 'Fideuá', 'Montaña'),
         'Tipo de paella', 'Elige un tipo de paella de la lista (con su tilde): el recargo se calcula a partir de él.')
    for cell, t in (('B9', 'Concepto premium'), ('B11', 'Show cooking'), ('B12', 'Suministra cliente')):
        c.dv(ws, cell, SINO, t, 'Elige «sí» o «no» de la lista (con tilde).')
    c.dv(ws, 'B13', lista('seca', 'lluvia'), 'Estación', 'Elige «seca» o «lluvia» de la lista.')

def financiero(wb, c):
    ws = wb['Proyección 3 años']
    for r in (5, 6):
        for col in 'BCD': c.fmt(ws, f'{col}{r}', CNT)

def proveedores(wb, c):
    ws = wb.worksheets[0]
    n = sum(1 for r in range(5, ws.max_row + 1) if ws.cell(r, 3).value)
    new = f'{n} proveedores'
    if ws.title != new:
        old = ws.title; c.log.append(f'hoja {old!r} -> {new!r}'); ws.title = new
        for w in wb.worksheets:   # referencias al nombre viejo
            for row in w.iter_rows():
                for cell in row:
                    if isinstance(cell.value, str) and old in cell.value and cell.value.startswith('='):
                        cell.value = cell.value.replace(f"'{old}'", f"'{new}'"); c.log.append(f'ref {w.title}!{cell.coordinate}')
                    if cell.hyperlink and cell.hyperlink.location and old in cell.hyperlink.location:
                        cell.hyperlink.location = cell.hyperlink.location.replace(old, new); c.log.append(f'link {cell.coordinate}')
    a1 = ws['A1'].value
    if isinstance(a1, str): c.val(ws, 'A1', re.sub(r'^\d+ Proveedores', f'{n} Proveedores', a1))
    return n

def checklist(wb, c):
    ws = wb.worksheets[0]
    items = [r for r in range(5, ws.max_row + 1) if isinstance(ws.cell(r, 2).value, int)]
    c.items = len(items); first, last = min(items), max(items)
    for r in items:
        if ws.cell(r, 6).value in (None, ''): c.val(ws, f'F{r}', '☐')
        ws.cell(r, 6).alignment = Alignment(horizontal='center', vertical='center')
    sq = f'F{first}:F{last}'
    c.dv(ws, sq, lista('✓', '☐', 'N/A'), 'Estado del hito', 'Elige ✓ (hecho), ☐ (pendiente) o N/A (no aplica).')
    a2 = ws['A2'].value
    if isinstance(a2, str): c.val(ws, 'A2', re.sub(r'\b\d+ hitos', f'{len(items)} hitos', a2))
    base = Font(name=ws['A2'].font.name, size=10)
    c.val(ws, 'A3', 'Marca ✓ al completar; N/A si no aplica')
    ws['A3'].font = Font(name=base.name, size=10, italic=True, color='595959')
    c.val(ws, 'D3', 'Progreso del checklist →'); ws['D3'].alignment = Alignment(horizontal='right')
    c.val(ws, 'G3', 'Completados'); c.val(ws, 'I3', 'Aplicables')
    c.val(ws, 'H3', f'=COUNTIF({sq},"✓")')
    c.val(ws, 'J3', f'=COUNTIF({sq},"✓")+COUNTIF({sq},"☐")')
    c.val(ws, 'E3', '=IFERROR(H3/J3,"")'); c.fmt(ws, 'E3', '0%')
    for k in ('D3', 'G3', 'I3'): ws[k].font = Font(name=base.name, size=10, bold=True)
    for k in ('E3', 'H3', 'J3'): ws[k].font = Font(name=base.name, size=11, bold=True); ws[k].alignment = Alignment(horizontal='center')
    ws['G3'].alignment = ws['I3'].alignment = Alignment(horizontal='right')
    if not any(sq in str(r.sqref) for r in ws.conditional_formatting):
        ws.conditional_formatting.add(sq, CellIsRule(operator='equal', formula=['"✓"'],
            fill=PatternFill(start_color='C6EFCE', end_color='C6EFCE', fill_type='solid'), font=Font(color='006100', bold=True)))
        c.log.append(f'formato condicional {sq}')

def common(wb, c):
    p = wb.properties
    if p.creator != 'AI Chef Pro': c.log.append(f'creator {p.creator!r}'); p.creator = 'AI Chef Pro'
    a1 = wb.worksheets[0]['A1'].value
    t = f'{a1} · v1.2'
    if p.title != t: c.log.append(f'title -> {t}'); p.title = t
    for ws in wb.worksheets:
        ps = ws.page_setup
        if ps.paperSize != 9: c.log.append(f'{ws.title} paperSize'); ps.paperSize = 9
        if ps.fitToWidth != 1 or ps.fitToHeight != 0 or not (ws.sheet_properties.pageSetUpPr and ws.sheet_properties.pageSetUpPr.fitToPage):
            pass  # ya estaba A4 + 1 página de ancho en los 4 (auditado); no se toca fitToHeight
        if not (ws.sheet_properties.pageSetUpPr and ws.sheet_properties.pageSetUpPr.fitToPage) or ps.fitToWidth != 1:
            ws.sheet_properties.pageSetUpPr = openpyxl.worksheet.properties.PageSetupProperties(fitToPage=True)
            ps.fitToWidth = 1; c.log.append(f'{ws.title} ajuste a 1 pág. de ancho')

JOBS = {
 'calculadora-pricing-eventos-paellero.xlsx': [calculadora],
 'plan-financiero-paellero-eventos.xlsx': [financiero],
 'plantilla-proveedores-paellero-eventos.xlsx': [proveedores],
 'checklist-apertura-paellero-eventos.xlsx': [checklist],
}
if __name__ == '__main__':
    total = 0; changed = []
    for f, fns in JOBS.items():
        wb = openpyxl.load_workbook(D + f); c = Ctx(f)
        for fn in fns: r = fn(wb, c)
        common(wb, c)
        extra = f' | items={getattr(c, "items", "")}' if hasattr(c, 'items') else ''
        print(f'== {f}: {len(c.log)} cambios{extra}')
        for l in c.log: print('   ', l)
        total += len(c.log)
        if c.log:
            wb.save(OUT + f); changed.append(OUT + f)
    if changed:
        sys.path.insert(0, f'{ROOT}/scripts/productos-digitales')
        import inject_cache
        for p in changed:
            nf, inj, fail = inject_cache.inject(p); print(f'  cache {os.path.basename(p)}: formulas={nf} inyectado={inj} fallos={fail}')
    print('TOTAL cambios:', total, '(APPLY)' if APPLY else '(dry-run en ' + OUT + ')')
