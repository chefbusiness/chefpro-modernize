import sys, re, glob, os, zipfile
from docx import Document
D = sys.argv[1]
OUT = sys.argv[2]; os.makedirs(OUT, exist_ok=True)
def iter_pars(doc):
    for p in doc.paragraphs: yield p
    def tbl(t):
        for r in t.rows:
            seen=set()
            for c in r.cells:
                if id(c._tc) in seen: continue
                seen.add(id(c._tc))
                for p in c.paragraphs: yield p
                for t2 in c.tables: yield from tbl(t2)
    for t in doc.tables: yield from tbl(t)
    for s in doc.sections:
        for p in s.header.paragraphs + s.footer.paragraphs: yield p
def text(path):
    return [p.text for p in iter_pars(Document(path))]
if __name__=='__main__':
    for f in sorted(glob.glob(D+'/*.docx')):
        b=os.path.basename(f)
        assert zipfile.ZipFile(f).testzip() is None
        d=Document(f); s=d.sections[0]
        pars=text(f)
        open(f'{OUT}/{b}.txt','w').write('\n'.join(pars))
        print(b, len(pars),'pars', len(d.tables),'tablas', round(s.page_width.cm,1),'x',round(s.page_height.cm,1), repr(d.core_properties.author), repr(d.core_properties.title), len(' '.join(pars).split()),'palabras')
