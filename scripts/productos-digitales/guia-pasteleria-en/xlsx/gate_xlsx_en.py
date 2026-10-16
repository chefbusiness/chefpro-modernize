#!/usr/bin/env python3
"""Gate de invariantes ES ↔ EN de los 8 xlsx de la guía Food Cost (SPEC D12). Corre en el VPS (pycel).

Por pareja de libros:
  G1 estructura: pestañas (mapa SPEC §2, mismo orden), mismas celdas con valor y del mismo tipo, celdas
     combinadas, paneles, anchos de columna, altos de fila, márgenes y ajuste de impresión, protección,
     formato condicional (nº y rangos) y validaciones (nº, rangos, tipo y fórmulas salvo las adaptaciones
     declaradas en config_en.py).
  G2 fórmulas: mismo nº por hoja y cada fórmula EN = fórmula ES token a token (openpyxl Tokenizer) tras mapear
     pestañas; los literales de texto se emparejan y deben formar un mapa coherente (un literal ES → un único EN).
  G3 valores (pycel): se evalúan ES, ES con los tipos de impuesto a 0 (config impuesto0) y EN. Toda fórmula EN
     debe valer lo mismo que ES-impuesto-0 (1e-6), o su texto traducido, o el mismo valor por un factor exacto
     de unidad (lb/kg, ml/cl, fl oz/cl) — que se lista. Las celdas que cambian de ES a ES-impuesto-0 son las
     que dependen del impuesto: se listan con su rótulo (deben ser impuestos y precios con impuesto).
  G4 residuo: cero español (residuo() de gate_md_en.py), cero escrituras no latinas, cero «€»/EUR/euro en
     textos, validaciones, pies, gráficos, docProps, formatos y XML crudo; formatos sin moneda, fechas mm/dd.
  G5 gráficos: los que tenga el ES existen en el EN (chart, drawing, rels, content types, <drawing> en la
     hoja) con título y series apuntando a la pestaña EN.
  --cache  G6: tras inject_cache.py, toda fórmula con valor tiene caché (data_only) igual al valor pycel.

    python3 gate_xlsx_en.py <dir_es> <dir_en> [--cache] [--json informe.json]
"""
import html
import json
import math
import os
import re
import sys
import zipfile
import logging
import contextlib

logging.disable(logging.CRITICAL)
AQUI = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, AQUI)
sys.path.insert(0, os.path.dirname(AQUI))
import openpyxl  # noqa: E402
from openpyxl.formula.tokenizer import Tokenizer, Token  # noqa: E402
from pycel import ExcelCompiler  # noqa: E402
import config_en as C  # noqa: E402
import gate_md_en as G  # noqa: E402

try:
    import numpy as np
except ImportError:  # pragma: no cover
    np = None

RX_MONEDA = re.compile(r'€|\bEUR\b|\beuros?\b', re.I)


def nativo(v):
    if np is not None and isinstance(v, np.generic):
        return v.item()
    return v


def expandir(rango):
    from aplicar_en import expandir as e
    return e(rango)


# ------------------------------------------------------------------------------------------- utilidades
def tokens(f):
    if f is None:
        return []
    f = f if f.startswith('=') else '=' + f
    return Tokenizer(f).items


def partir_hoja(v):
    if '!' not in v:
        return None, v
    h, r = v.rsplit('!', 1)
    if h.startswith("'"):
        h = h[1:-1].replace("''", "'")
    return h, r


def comparar_formula(fes, fen, hojas, pares, donde, errores):
    tes, ten = tokens(fes), tokens(fen)
    if len(tes) != len(ten):
        errores.append('%s: nº de tokens %d ≠ %d\n    ES %s\n    EN %s' % (donde, len(tes), len(ten), fes, fen))
        return
    for a, b in zip(tes, ten):
        if (a.type, a.subtype) != (b.type, b.subtype):
            errores.append('%s: token %r ≠ %r' % (donde, a.value, b.value))
            return
        if a.type == Token.OPERAND and a.subtype == Token.TEXT:
            la, lb = a.value[1:-1].replace('""', '"'), b.value[1:-1].replace('""', '"')
            pares.setdefault(la, set()).add(lb)
        elif a.type == Token.OPERAND and a.subtype == Token.RANGE:
            ha, ra = partir_hoja(a.value)
            hb, rb = partir_hoja(b.value)
            if ra != rb or (ha is None) != (hb is None) or (ha is not None and hojas.get(ha) != hb):
                errores.append('%s: referencia %s ≠ %s' % (donde, a.value, b.value))
        elif a.value != b.value:
            errores.append('%s: token %r ≠ %r' % (donde, a.value, b.value))


def igual(a, b):
    a, b = nativo(a), nativo(b)
    if isinstance(a, bool) or isinstance(b, bool):
        return a == b
    if isinstance(a, (int, float)) and isinstance(b, (int, float)):
        return math.isclose(a, b, rel_tol=1e-9, abs_tol=1e-6)
    if a in (None, '') and b in (None, ''):
        return True
    return a == b


def factor_unidad(a, b):
    a, b = nativo(a), nativo(b)
    if not (isinstance(a, (int, float)) and isinstance(b, (int, float))) or isinstance(a, bool) or a == 0:
        return None
    r = b / a
    for nombre, k in C.FACTORES_UNIDAD.items():
        if math.isclose(r, k, rel_tol=1e-9):
            return nombre
    return None


def rotulo(ws, ref):
    """Encabezado de columna (la fila de cabecera más cercana por encima) + rótulo de la fila."""
    c = ws[ref]
    fila = None
    for col in range(1, c.column):
        v = ws.cell(c.row, col).value
        if isinstance(v, str) and not str(v).startswith('='):
            fila = v
    cab = None
    for r in range(c.row - 1, 0, -1):
        v = ws.cell(r, c.column).value
        if isinstance(v, str) and not v.startswith('='):
            cab = v
            break
    return ' | '.join(x[:60] for x in (fila, cab) if x)


def estilo_de(c):
    nf = c.number_format or 'General'
    nf = C.FORMATOS.get(nf, nf)
    return (repr(c.font), repr(c.fill), repr(c.border), repr(c.alignment), repr(c.protection), nf)


def textos_residuo(txt, donde, errores, avisos):
    if not isinstance(txt, str) or not txt.strip():
        return
    e, a = G.residuo(txt)
    for x in e:
        errores.append('%s: %s · %r' % (donde, x, txt[:90]))
    for x in a:
        avisos.append('%s: %s' % (donde, x))
    nl = G.no_latinos(txt)
    if nl:
        errores.append('%s: escritura no latina %s' % (donde, nl))
    if RX_MONEDA.search(txt):
        errores.append('%s: moneda € / EUR / euro · %r' % (donde, txt[:90]))


# ------------------------------------------------------------------------------------------- G1-G2, G4-G5
def estructura(libro, cfg, wes, wen, errores, avisos, pares):
    hojas = cfg['hojas']
    if wen.sheetnames != [hojas[s] for s in wes.sheetnames]:
        errores.append('pestañas %s ≠ %s' % (wen.sheetnames, [hojas[s] for s in wes.sheetnames]))
        return {}
    textos_residuo(wen.properties.title, 'docProps title', errores, avisos)
    textos_residuo(wen.properties.subject, 'docProps subject', errores, avisos)
    if wen.properties.creator != 'AI Chef Pro':
        errores.append('docProps creator %r' % wen.properties.creator)
    dv_cfg = {(h, r): v for (h, r), v in cfg.get('dv', {}).items()}
    nform = {}
    for es_ws, en_ws in zip(wes.worksheets, wen.worksheets):
        h = es_ws.title
        textos_residuo(en_ws.title, 'pestaña', errores, avisos)
        ce = {c.coordinate: c for row in es_ws.iter_rows() for c in row if c.value is not None}
        cn = {c.coordinate: c for row in en_ws.iter_rows() for c in row if c.value is not None}
        if set(ce) != set(cn):
            errores.append('%s: celdas con valor distintas: solo ES %s · solo EN %s' % (
                h, sorted(set(ce) - set(cn))[:10], sorted(set(cn) - set(ce))[:10]))
        nf_es = nf_en = 0
        for ref in sorted(set(ce) & set(cn)):
            a, b = ce[ref], cn[ref]
            fa, fb = a.data_type == 'f', b.data_type == 'f'
            nf_es += fa
            nf_en += fb
            if fa != fb:
                errores.append('%s!%s: fórmula en un lado y no en el otro' % (h, ref))
                continue
            if fa:
                comparar_formula(a.value, b.value, hojas, pares, '%s!%s' % (h, ref), errores)
            elif isinstance(a.value, str) != isinstance(b.value, str):
                errores.append('%s!%s: tipo %s ≠ %s' % (h, ref, type(a.value).__name__, type(b.value).__name__))
            if isinstance(b.value, str) and not fb:
                textos_residuo(b.value, '%s!%s' % (en_ws.title, ref), errores, avisos)
            if estilo_de(a) != estilo_de(b) and not any(h == hh and ref in expandir(rr)
                                                        for hh, rr, _ in cfg.get('estilo', [])):
                errores.append('%s!%s: estilo distinto %s ≠ %s' % (h, ref, estilo_de(a)[-1], estilo_de(b)[-1]))
            if b.number_format and RX_MONEDA.search(b.number_format):
                errores.append('%s!%s: formato con moneda %r' % (h, ref, b.number_format))
            if 'dd/mm' in (b.number_format or ''):
                errores.append('%s!%s: fecha dd/mm %r' % (h, ref, b.number_format))
        nform[en_ws.title] = (nf_es, nf_en)
        if nf_es != nf_en:
            errores.append('%s: nº de fórmulas %d ≠ %d' % (h, nf_es, nf_en))
        # rejilla, paneles, impresión, protección
        for nombre, va, vb in (
                ('combinadas', sorted(map(str, es_ws.merged_cells.ranges)), sorted(map(str, en_ws.merged_cells.ranges))),
                ('paneles', es_ws.freeze_panes, en_ws.freeze_panes),
                ('anchos', {k: v.width for k, v in es_ws.column_dimensions.items()},
                 {k: v.width for k, v in en_ws.column_dimensions.items()}),
                ('altos', {k: v.height for k, v in es_ws.row_dimensions.items() if v.height},
                 {k: v.height for k, v in en_ws.row_dimensions.items() if v.height}),
                ('orientación', es_ws.page_setup.orientation, en_ws.page_setup.orientation),
                ('ajuste', (es_ws.page_setup.fitToWidth, es_ws.page_setup.fitToHeight),
                 (en_ws.page_setup.fitToWidth, en_ws.page_setup.fitToHeight)),
                ('márgenes', (es_ws.page_margins.left, es_ws.page_margins.top),
                 (en_ws.page_margins.left, en_ws.page_margins.top)),
                ('protección', (es_ws.protection.sheet, es_ws.protection.password),
                 (en_ws.protection.sheet, en_ws.protection.password)),
                ('autofiltro', es_ws.auto_filter.ref, en_ws.auto_filter.ref)):
            if va != vb:
                errores.append('%s: %s %s ≠ %s' % (h, nombre, va, vb))
        if str(en_ws.page_setup.paperSize) != C.PAPEL_LETTER:
            errores.append('%s: papel %s (esperado US Letter)' % (h, en_ws.page_setup.paperSize))
        for hf in (en_ws.oddFooter, en_ws.oddHeader, en_ws.evenFooter, en_ws.firstFooter):
            for parte in (hf.left, hf.center, hf.right):
                textos_residuo(parte.text, '%s pie' % en_ws.title, errores, avisos)
        # formato condicional
        cfa = [(str(cf.sqref), r.type, r.operator) for cf in es_ws.conditional_formatting for r in cf.rules]
        cfb = [(str(cf.sqref), r.type, r.operator) for cf in en_ws.conditional_formatting for r in cf.rules]
        if cfa != cfb:
            errores.append('%s: formato condicional distinto' % h)
        else:
            ra = [r for cf in es_ws.conditional_formatting for r in cf.rules]
            rb = [r for cf in en_ws.conditional_formatting for r in cf.rules]
            for x, y in zip(ra, rb):
                for fx, fy in zip(x.formula or [], y.formula or []):
                    comparar_formula(fx, fy, hojas, pares, '%s CF' % h, errores)
        # validaciones
        dva = {str(d.sqref): d for d in es_ws.data_validations.dataValidation}
        dvb = {str(d.sqref): d for d in en_ws.data_validations.dataValidation}
        for sq, d in dva.items():
            primera = sq.split()[0].split(':')[0]
            o = dv_cfg.get((h, primera), 'nada')
            if o is None:
                if sq in dvb:
                    errores.append('%s DV %s: debía eliminarse' % (h, sq))
                else:
                    avisos.append('%s DV %s: eliminada (adaptación declarada)' % (h, sq))
                continue
            if sq not in dvb:
                errores.append('%s DV %s: falta en el EN' % (h, sq))
                continue
            e = dvb[sq]
            for k in ('error', 'errorTitle', 'prompt', 'promptTitle'):
                textos_residuo(getattr(e, k), '%s DV %s %s' % (en_ws.title, primera, k), errores, avisos)
                if bool(getattr(d, k)) != bool(getattr(e, k)):
                    errores.append('%s DV %s: %s presente en un solo lado' % (h, primera, k))
            for k in ('showDropDown', 'showErrorMessage', 'showInputMessage', 'allowBlank'):
                if getattr(d, k) != getattr(e, k):
                    errores.append('%s DV %s: %s distinto' % (h, primera, k))
            if o != 'nada':
                avisos.append('%s DV %s: adaptada (%s → %s %s %s)' % (h, primera, d.type, e.type, e.formula1,
                                                                    e.formula2 or ''))
                if e.type == 'list':
                    for it in e.formula1.strip('"').split(','):
                        textos_residuo(it, '%s DV lista' % en_ws.title, errores, avisos)
                continue
            if (d.type, d.operator, d.formula2) != (e.type, e.operator, e.formula2):
                errores.append('%s DV %s: tipo/operador/fórmula2 distintos' % (h, primera))
            if d.type == 'list' and d.formula1 and d.formula1.startswith('"'):
                ia, ib = d.formula1.strip('"').split(','), e.formula1.strip('"').split(',')
                if len(ia) != len(ib):
                    errores.append('%s DV %s: lista de %d ≠ %d' % (h, primera, len(ia), len(ib)))
                for x, y in zip(ia, ib):
                    pares.setdefault(x, set()).add(y)
                    textos_residuo(y, '%s DV lista' % en_ws.title, errores, avisos)
            elif d.formula1 != e.formula1:
                comparar_formula(d.formula1, e.formula1, hojas, pares, '%s DV %s' % (h, primera), errores)
        for sq in dvb:
            if sq not in dva:
                errores.append('%s DV %s: solo en el EN' % (h, sq))
    # nombres definidos (títulos de impresión, autofiltros)
    na = sorted((d.name, d.localSheetId) for ws in wes.worksheets for d in ws.defined_names.values()) \
        if hasattr(wes.worksheets[0], 'defined_names') else []
    nb = sorted((d.name, d.localSheetId) for ws in wen.worksheets for d in ws.defined_names.values()) \
        if hasattr(wen.worksheets[0], 'defined_names') else []
    if na != nb:
        errores.append('nombres definidos %s ≠ %s' % (na, nb))
    for es_ws, en_ws in zip(wes.worksheets, wen.worksheets):
        if es_ws.print_title_rows != en_ws.print_title_rows:
            errores.append('%s: títulos de impresión %s ≠ %s' % (es_ws.title, es_ws.print_title_rows,
                                                                 en_ws.print_title_rows))
    return nform


def paquete(ruta_es, ruta_en, hojas, errores, avisos):
    za, zb = zipfile.ZipFile(ruta_es), zipfile.ZipFile(ruta_en)
    na, nb = za.namelist(), zb.namelist()
    if na != nb:
        errores.append('partes del paquete distintas: %s' % sorted(set(na) ^ set(nb)))
    graficos = [n for n in na if n.startswith('xl/charts/')]
    import xml.etree.ElementTree as ET
    for n in nb:
        if n.endswith('.xml') or n.endswith('.rels'):
            try:
                ET.fromstring(zb.read(n))
            except ET.ParseError as e:
                errores.append('%s: XML mal formado (%s)' % (n, e))
        x = zb.read(n).decode('utf-8', 'replace')
        crudo = html.unescape(x)
        if RX_MONEDA.search(crudo):
            errores.append('%s: «€»/EUR/euro en el XML: %r' % (n, RX_MONEDA.search(crudo).group(0)))
        if n == 'xl/styles.xml':
            for code in re.findall(r'formatCode="([^"]*)"', x):
                code = html.unescape(code)
                if 'dd/mm' in code or RX_MONEDA.search(code) or '[$' in code:
                    errores.append('formato %r' % code)
    for n in graficos:
        if n not in nb:
            errores.append('falta el gráfico %s' % n)
            continue
        x = html.unescape(zb.read(n).decode('utf-8'))
        for t in re.findall(r'<a:t>(.*?)</a:t>', x, re.S):
            textos_residuo(t, n, errores, avisos)
        for f in re.findall(r'<(?:c:)?f>(.*?)</(?:c:)?f>', x, re.S):
            h, _ = partir_hoja(f.split(':')[0])
            if h not in hojas.values():
                errores.append('%s: referencia a pestaña no EN %r' % (n, f))
    if graficos:
        ct = zb.read('[Content_Types].xml').decode()
        for n in graficos:
            if '/' + n not in ct:
                errores.append('content types sin %s' % n)
        dibujos = [n for n in nb if re.match(r'xl/worksheets/sheet\d+\.xml$', n) and '<drawing ' in
                   zb.read(n).decode()]
        if not dibujos:
            errores.append('ninguna hoja EN referencia el dibujo del gráfico')
    return len(graficos)


# ------------------------------------------------------------------------------------------- G3 valores
def evaluar(xl, celdas):
    out = {}
    for h, ref in celdas:
        try:
            with open(os.devnull, 'w') as dn, contextlib.redirect_stderr(dn):
                out[(h, ref)] = nativo(xl.evaluate("'%s'!%s" % (h.replace("'", "''"), ref)))
        except Exception as e:  # noqa: BLE001
            out[(h, ref)] = 'ERROR pycel: %s' % type(e).__name__
    return out


def valores(ruta_es, ruta_en, cfg, wes, wen, pares, errores, avisos):
    hojas = cfg['hojas']
    celdas = [(ws.title, c.coordinate) for ws in wes.worksheets for row in ws.iter_rows() for c in row
              if c.data_type == 'f']
    # pycel devuelve el valor cacheado de una fórmula si lo hay: se evalúan COPIAS sin caché
    ruta_es0 = crear_t0(ruta_es, None, wes, avisos)
    ruta_t0 = crear_t0(ruta_es, cfg, wes, avisos)
    ruta_en0 = crear_t0(ruta_en, None, wen, avisos)
    xl_es = ExcelCompiler(filename=ruta_es0)
    xl_t0 = ExcelCompiler(filename=ruta_t0)
    xl_en = ExcelCompiler(filename=ruta_en0)
    v_es = evaluar(xl_es, celdas)
    v_t0 = evaluar(xl_t0, celdas)
    v_en = evaluar(xl_en, [(hojas[h], r) for h, r in celdas])
    for r in (ruta_es0, ruta_t0, ruta_en0):
        os.remove(r)
    lit = {}
    for a, b in pares.items():
        lit.setdefault(a, set()).update(b)
    for es_ws, en_ws in zip(wes.worksheets, wen.worksheets):        # textos de celda (los leen las fórmulas)
        for row in es_ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.data_type != 'f':
                    d = en_ws[c.coordinate].value
                    if isinstance(d, str):
                        lit.setdefault(c.value, set()).add(d)
    impuesto, unidad, fallos = [], [], 0
    for h, ref in celdas:
        a, t, b = v_es[(h, ref)], v_t0[(h, ref)], v_en[(hojas[h], ref)]
        for x in (a, t, b):
            if isinstance(x, str) and x.startswith('ERROR pycel'):
                errores.append('%s!%s: %s' % (h, ref, x))
        etiqueta = rotulo(wen[hojas[h]], ref)
        if not igual(a, t):
            impuesto.append('%s!%s [%s] ES %s → impuesto 0 %s' % (hojas[h], ref, etiqueta, fmt(a), fmt(t)))
        if igual(t, b):
            continue
        if isinstance(t, str) and isinstance(b, str) and b in lit.get(t, {t}):
            continue
        k = factor_unidad(t, b)
        if k:
            unidad.append('%s!%s [%s] ×%s' % (hojas[h], ref, etiqueta, k))
            continue
        fallos += 1
        errores.append('%s!%s [%s]: EN %s ≠ ES-impuesto-0 %s' % (hojas[h], ref, etiqueta, fmt(b), fmt(t)))
    return len(celdas), impuesto, unidad, v_en


def crear_t0(ruta_es, cfg, wes, avisos):
    """Copia del ES con los tipos de impuesto a 0 (y el precio del menú del día en neto): la referencia del EN."""
    import tempfile
    cambios = {}
    for item in (cfg or {}).get('impuesto0', []):
        h, rango = item[0], item[1]
        v = item[2] if len(item) > 2 else 0
        for ref in expandir(rango):
            if isinstance(wes[h][ref].value, (int, float)):
                cambios.setdefault(h, {})[ref] = v
    z = zipfile.ZipFile(ruta_es)
    wb = z.read('xl/workbook.xml').decode()
    rels = z.read('xl/_rels/workbook.xml.rels').decode()
    relmap = {re.search(r'Id="([^"]+)"', r).group(1): re.search(r'Target="([^"]+)"', r).group(1)
              for r in re.findall(r'<Relationship [^>]*/>', rels)}
    parte = {}
    for mm in re.finditer(r'<sheet [^>]*name="([^"]+)"[^>]*r:id="([^"]+)"', wb):
        tgt = relmap[mm.group(2)]
        parte[html.unescape(mm.group(1))] = tgt.lstrip('/') if tgt.startswith('/') else 'xl/' + tgt
    fd, ruta = tempfile.mkstemp(suffix='.xlsx', prefix='t0_')
    os.close(fd)
    n = 0
    with zipfile.ZipFile(ruta, 'w', zipfile.ZIP_DEFLATED) as zo:
        for info in z.infolist():
            d = z.read(info.filename)
            if info.filename.startswith('xl/worksheets/sheet'):
                d = re.sub(rb'(<c r="[A-Z]+\d+")([^>]*>)(<f>[^<]*</f>)(?:<v>[^<]*</v>|<v\s*/>)?',
                           lambda m: m.group(1) + re.sub(rb'\s+t="[^"]*"', b'', m.group(2)) + m.group(3), d)
            for h, refs in cambios.items():
                if parte[h] == info.filename:
                    x = d.decode()
                    for ref, v in refs.items():
                        x, k = re.subn(r'(<c r="%s"[^>]*>)<v>[^<]*</v>' % ref, r'\g<1><v>%r</v>' % float(v), x)
                        n += k
                    d = x.encode()
            zo.writestr(info, d)
    if cfg is not None:
        avisos.append('referencia ES-impuesto-0: %d entradas puestas a 0 (o a neto)' % n)
    return ruta


def fmt(v):
    return ('%.6g' % v) if isinstance(v, float) else repr(v)


# ------------------------------------------------------------------------------------------- G6 caché
def cache(ruta_en, wen_titles, v_en, errores):
    wd = openpyxl.load_workbook(ruta_en, data_only=True)
    con, vacias = 0, []
    for (h, ref), v in v_en.items():
        c = wd[h][ref].value
        if hasattr(c, 'year'):                       # celda con formato de fecha: openpyxl da datetime
            from openpyxl.utils.datetime import to_excel
            c = to_excel(c)
        if v in (None, ''):
            continue
        if c is None:
            vacias.append('%s!%s' % (h, ref))
        elif not igual(c, v):
            errores.append('caché %s!%s %r ≠ pycel %r' % (h, ref, c, v))
        else:
            con += 1
    for x in vacias:
        errores.append('fórmula con valor sin caché: %s' % x)
    return con, len(vacias)


# ------------------------------------------------------------------------------------------- main
def main():
    argv = sys.argv[1:]
    if '--json' in argv:
        i = argv.index('--json')
        del argv[i:i + 2]
    args = [a for a in argv if not a.startswith('--')]
    dir_es, dir_en = args[0], args[1]
    con_cache = '--cache' in sys.argv
    ruta_json = sys.argv[sys.argv.index('--json') + 1] if '--json' in sys.argv else None
    libros = args[2:] or list(C.LIBROS)
    total_err, informe = 0, {}
    for libro in libros:
        cfg = C.LIBROS[libro]
        ruta_es, ruta_en = os.path.join(dir_es, libro), os.path.join(dir_en, cfg['out'])
        errores, avisos, pares = [], [], {}
        wes, wen = openpyxl.load_workbook(ruta_es), openpyxl.load_workbook(ruta_en)
        nform = estructura(libro, cfg, wes, wen, errores, avisos, pares)
        ngraf = paquete(ruta_es, ruta_en, cfg['hojas'], errores, avisos)
        for a, b in pares.items():
            if len(b) > 1:
                errores.append('literal ES %r traducido de %d formas: %s' % (a, len(b), sorted(b)))
            for y in b:
                textos_residuo(y, 'literal de fórmula/CF/DV', errores, avisos)
        n, impuesto, unidad, v_en = valores(ruta_es, ruta_en, cfg, wes, wen, pares, errores, avisos)
        cache_ok = None
        if con_cache:
            cache_ok = cache(ruta_en, wen.sheetnames, v_en, errores)
        total_err += len(errores)
        print('=' * 100)
        print('%s ↔ %s · %s' % (libro, cfg['out'], 'OK' if not errores else '%d ERRORES' % len(errores)))
        print('  fórmulas por hoja (ES/EN): %s' % ', '.join('%s %d/%d' % (k, a, b) for k, (a, b) in nform.items()))
        print('  fórmulas evaluadas con pycel: %d · literales emparejados: %d · gráficos: %d' % (n, len(pares), ngraf))
        print('  dependen del impuesto (ES ≠ ES con impuesto 0; el EN iguala al segundo): %d' % len(impuesto))
        for x in impuesto:
            print('     ', x)
        print('  escaladas por unidad (EN = ES × factor exacto): %d' % len(unidad))
        for x in unidad:
            print('     ', x)
        if cache_ok:
            print('  caché data_only verificada: %d celdas con valor, %d sin caché' % cache_ok)
        for x in sorted(set(avisos)):
            print('  AVISO', x)
        for x in errores[:60]:
            print('  ERROR', x)
        informe[libro] = {'out': cfg['out'], 'errores': errores, 'avisos': sorted(set(avisos)),
                          'impuesto': impuesto, 'unidad': unidad, 'formulas': nform, 'graficos': ngraf}
    if ruta_json:
        json.dump(informe, open(ruta_json, 'w'), ensure_ascii=False, indent=1)
    print('=' * 100)
    print('TOTAL ERRORES: %d' % total_err)
    sys.exit(1 if total_err else 0)


if __name__ == '__main__':
    main()
