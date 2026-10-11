#!/usr/bin/env python3
"""
docx_pae_es.py — Paellero ES 2.2 (SPEC D11; copia de `catering-es-2_0/docx_cat_es.py`): ensambla
`plan-de-negocio-paellero-eventos.docx` en las 131 ranuras del docx del Food Truck ES (10 Heading 1), con los textos de
`textos/pae_es_{a,b}.json` y las cifras de `cifras_caso_es.json` (tokens {ASI}, leídos de la caché del Excel 2.2 por
`cifras_pae_es.py`). Fija portada, aviso, índice, encabezados y cierre, y AÑADE el anexo autonómico del fuego (11.º
Heading 1), trasplantado de la 1.2 (p93-p110 en `git show 2d6cac44`) con las correcciones del research (SPEC D8: la
licencia de una finca no levanta la restricción en la franja forestal; sin «excepción cultural»). A4; autor
`AI Chef Pro`. Gates al final: 11 Heading 1, A4, 0 tokens sin resolver, 0 cifras/palabras prohibidas, 0 no latinos,
≥ 6.000 palabras. Corre en el VPS (python-docx).

    /root/venv-guias/bin/python docx_pae_es.py --salida <dir>          # escribe <dir>/plan-de-negocio-…docx
    /root/venv-guias/bin/python docx_pae_es.py --salida <dir> --parcial # ranuras sin texto: quedan con «[PENDIENTE]»

Sesión Claude Code, 11-oct-2026.
"""
import argparse
import datetime
import json
import os
import re
import shutil
import sys

import docx
from docx.shared import Mm

AQUI = os.path.dirname(os.path.abspath(__file__))
REPO = os.path.abspath(os.path.join(AQUI, '..', '..', '..'))
FT = os.path.join(REPO, 'astro-site', 'public', 'dl', 'plan-negocio-food-truck', 'plan-de-negocio-food-truck.docx')
NOMBRE = 'plan-de-negocio-paellero-eventos.docx'
sys.path.insert(0, AQUI)
sys.path.insert(0, os.path.join(os.path.dirname(AQUI), 'planes-v2_0'))
import motor                                                                       # noqa: E402
import tandas_docx as T                                                            # noqa: E402

FIJOS = {
    6: 'PLAN DE NEGOCIO',
    7: 'Paellero / Paella para Eventos',
    9: 'Plan financiero completo, análisis de mercado,\nestrategia operativa y checklist de apertura\npara montar un '
       'servicio de paella y arroces en directo para eventos',
    14: 'ChefBusiness Consultoría Gastronómica',
    16: 'Versión %s · %s' % (motor.VERSION, motor.VERSION_MES),   # la misma que estampa el motor en los xlsx
    18: 'AVISO LEGAL',
    19: 'Este documento ha sido elaborado por ChefBusiness Consultoría Gastronómica con fines informativos y '
        'orientativos. Las cifras salen del Excel del plan financiero que lo acompaña, con un caso de ejemplo que '
        'debes sustituir por el tuyo; las fuentes de los datos de mercado y de las normas se citan en el texto. Los '
        'resultados reales pueden variar. No constituye asesoramiento financiero, legal ni fiscal.',
    22: 'ÍNDICE',
    35: 'Anexos: Plan Financiero Excel + Checklist de Apertura Excel + Calculadora de Precios Excel',
    124: 'Consultoría Gastronómica',
    127: '→ Kit de Tareas: Catering / Eventos',
    128: '→ Kit de Escandallos Pro',
    129: '→ Pack Plantillas APPCC',
}
ENCABEZADOS = dict(zip((37, 44, 51, 59, 67, 76, 83, 91, 100, 109), (T.SECCIONES[i] for i in range(1, 11))))
INDICE = dict(zip(range(24, 34), ('%d.  %s' % (i, T.SECCIONES[i].split('. ', 1)[1]) for i in range(1, 11))))


#: Anexo autonómico (SPEC D3/D8/D11): texto de la 1.2 corregido con `F1-research-paellero.md` §1 y §5.
ANEXO_TITULO = 'Anexo: uso del fuego al aire libre por comunidad autónoma'
ANEXO = [
    (None, 'El uso del fuego al aire libre lo regula cada comunidad autónoma en los montes y en sus zonas colindantes '
           '(Ley 43/2003, de Montes, art. 44.3): suele haber una franja de 400 a 500 metros alrededor del terreno '
           'forestal y una época de peligro alto en verano, y en los niveles de alerta máximos quedan sin efecto las '
           'autorizaciones; algunas alertas prohíben también el gas. Si la finca del evento está en terreno forestal o '
           'en esa franja, la restricción se aplica aunque la finca tenga licencia; fuera de ella, manda la ordenanza '
           'municipal. Esta es una guía de partida para las ocho comunidades con más demanda: confirma siempre con el '
           'ayuntamiento y con la comunidad antes de cada evento, y lleva en el contrato la cláusula de fuerza mayor y '
           'el plan B.'),
    ('Madrid', 'Plan INFOMA, con época de peligro desde el 15 de junio. Autorización municipal por evento: pídela con '
               '30-45 días de antelación. En una finca privada dentro de la franja forestal se aplican las mismas '
               'restricciones. Aporta el plan APPCC, el seguro de responsabilidad civil y la formación en higiene '
               'alimentaria del equipo.'),
    ('Cataluña', "Pla INFOCAT (Pla especial d'emergències per incendis forestals de Catalunya). En los niveles de "
                 'alerta más altos se prohíbe el fuego al aire libre en las zonas forestales. Autorización por evento '
                 "del ajuntament. En una finca privada fuera de la franja forestal, permiso del propietario y "
                 'protocolo de extinción. La documentación municipal puede estar en catalán.'),
    ('Baleares', 'Plan INFOBAL. Restricciones más severas de julio a septiembre por el riesgo de incendio. Autorización '
                 "del ajuntament y, en terreno forestal, de la conselleria. En finca privada: extintor homologado y "
                 'cubeto de retención; si está en la franja forestal, las mismas restricciones.'),
    ('Andalucía', 'Plan INFOCA. Época de peligro alto del 1 de junio al 15 de octubre, con una franja de 400 metros '
                  'alrededor del terreno forestal en la que no se puede usar fuego para cocinar; la licencia de la '
                  'finca no levanta esa restricción. Fuera de ella, autorización del ayuntamiento con unos 30 días de '
                  'antelación.'),
    ('Comunidad Valenciana', "Pla especial davant del risc d'incendis forestals de la Comunitat Valenciana. El uso "
                             'festivo-recreativo del fuego en suelo forestal o a menos de 500 metros necesita '
                             'declaración responsable y autorización del emplazamiento (arts. 145-146 del Decreto '
                             '98/1995, modificados por el Decreto 148/2018), y en preemergencia de nivel 3 quedan sin '
                             "efecto. En la vía pública, autorización del ajuntament: las paellas de Fallas y de las "
                             'fiestas patronales se rigen por el bando de cada municipio.'),
    ('País Vasco', 'Plan Especial de Emergencias por Riesgo de Incendios Forestales de Euskadi. Restricciones por '
                   'temporada en zonas forestales. Autorización del ayuntamiento; fuera de la franja forestal, manda '
                   'la ordenanza municipal. La documentación puede estar en euskera y en castellano.'),
    ('Galicia', 'PLADIGA (Plan de prevención e defensa contra os incendios forestais de Galicia). Restricciones '
                'marcadas de junio a septiembre. Autorización del concello; fuera de la franja forestal, manda la '
                'ordenanza municipal. La documentación puede estar en gallego y en castellano.'),
    ('Castilla y León', 'Plan INFOCAL. Época de peligro alto del 12 de junio al 12 de octubre en 2026, con una franja '
                        'de 400 metros; prohibido usar fuego para cocinar en zona forestal, incluso en barbacoas, y la '
                        'licencia de la finca no levanta la restricción. Fuera de ella, autorización del '
                        'ayuntamiento.'),
]


def anexo(d):
    """Inserta el anexo tras el §10 (ranura 117) y antes del bloque de cierre (p118-p130 del molde): salto de página,
    Heading 1 y un párrafo por comunidad (nombre en negrita). Se llama DESPUÉS de escribir las ranuras por índice."""
    from docx.enum.text import WD_BREAK
    ancla = d.paragraphs[118]
    ancla.insert_paragraph_before('').add_run().add_break(WD_BREAK.PAGE)
    ancla.insert_paragraph_before(ANEXO_TITULO, style='Heading 1')
    for nombre, texto in ANEXO:
        p = ancla.insert_paragraph_before('', style='Normal')
        if nombre:
            p.add_run(nombre + '. ').bold = True
        p.add_run(texto)


def poner(p, texto):
    if not p.runs:
        p.add_run(texto)
        return
    p.runs[0].text = texto
    for r in p.runs[1:]:
        r.text = ''


def ensamblar(salida, parcial=False):
    cif = json.load(open(os.path.join(AQUI, 'cifras_caso_es.json'), encoding='utf-8'))['tokens']
    textos = {}
    for t in ('pae_es_a', 'pae_es_b'):
        p = os.path.join(T.TEXTOS, t + '.json')
        if os.path.exists(p):
            textos.update(json.load(open(p, encoding='utf-8')))
    destino = os.path.join(salida, NOMBRE)
    shutil.copy2(FT, destino)
    d = docx.Document(destino)
    ps = d.paragraphs
    assert len(ps) == 131, len(ps)
    pendientes = []
    for r in T.ranuras():
        i = r['i']
        t = textos.get(str(i))
        if t is None:
            if not parcial:
                raise SystemExit('falta la ranura %d (usa --parcial para un ensayo)' % i)
            pendientes.append(i)
            t = '[PENDIENTE] ' + T.SECCIONES[r['seccion']]
        t = re.sub(r'\{([A-Z0-9_]+)\}', lambda m: cif.get(m.group(1), m.group(0)), t)
        poner(ps[i], t)
    for i, t in list(FIJOS.items()) + list(ENCABEZADOS.items()) + list(INDICE.items()):
        poner(ps[i], t)
    for s in d.sections:
        s.page_width, s.page_height = Mm(210), Mm(297)
    cp = d.core_properties
    cp.author, cp.last_modified_by = 'AI Chef Pro', 'AI Chef Pro'
    cp.title = 'Plan de negocio — Paellero / Paella para Eventos · v' + motor.VERSION
    cp.subject = 'Plan de Negocio: Paellero / Paella para Eventos'
    cp.comments = 'aichef.pro/plan-negocio-paellero-eventos'   # sin «generated by python-docx» (catering m13)
    anexo(d)
    cp.created = cp.modified = datetime.datetime.now().replace(microsecond=0)
    cp.revision = 1
    d.save(destino)
    return destino, pendientes


def gates(path, parcial):
    d = docx.Document(path)
    texto = '\n'.join(p.text for p in d.paragraphs)
    err = []
    h1 = [p.text for p in d.paragraphs if p.style.name == 'Heading 1']
    if len(h1) != 11:
        err.append('Heading 1: %d (esperados 11: 10 secciones + anexo)' % len(h1))
    s = d.sections[0]
    if abs(s.page_width - Mm(210)) > Mm(1) or abs(s.page_height - Mm(297)) > Mm(1):
        err.append('no es A4')
    sin = re.findall(r'\{[A-Z0-9_]+\}', texto)
    if sin:
        err.append('tokens sin resolver: %s' % sorted(set(sin)))
    for rx in T.PROHIBIDOS:
        m = re.search(rx, texto, re.I)
        if m:
            err.append('prohibido /%s/: …%s…' % (rx, texto[max(0, m.start() - 40):m.end() + 40].replace('\n', ' ')))
    if re.search(r'[Ͱ-ϿЀ-ӿ֐-ۿ฀-๿぀-ヿ㐀-鿿가-힯]', texto):
        err.append('caracteres no latinos')
    palabras = len(texto.split())
    if palabras < 6000 and not parcial:
        err.append('%d palabras (< 6.000)' % palabras)
    for x in err:
        print('ERROR', x)
    print('docx: %d párrafos · %d Heading 1 · %d palabras · %s' % (
        len(d.paragraphs), len(h1), palabras, 'VERDE' if not err else 'ROJO'))
    return 1 if err else 0


if __name__ == '__main__':
    ap = argparse.ArgumentParser()
    ap.add_argument('--salida', required=True)
    ap.add_argument('--parcial', action='store_true')
    a = ap.parse_args()
    os.makedirs(a.salida, exist_ok=True)
    path, pend = ensamblar(a.salida, a.parcial)
    print('→', path, '(%d ranuras pendientes)' % len(pend) if pend else '')
    sys.exit(gates(path, a.parcial))
