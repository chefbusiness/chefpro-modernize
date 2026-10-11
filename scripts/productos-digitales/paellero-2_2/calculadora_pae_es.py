#!/usr/bin/env python3
"""
calculadora_pae_es.py — Plan de Negocio Paellero para Eventos ES 2.2 (SPEC D12): `calculadora-pricing-eventos-paellero.xlsx`,
del coste de la paella al precio por invitado y al presupuesto completo. Fork de `catering-es-2_0/calculadora_cat_es.py`:
  · tabla de escalado por diámetro de paellera (manual técnico 1.2, tabla 1): raciones, arroz y caldo por paellera;
  · paelleras = invitados / raciones por paellera (hacia arriba); paelleros = paelleras / 3 (hacia arriba);
  · IVA del servicio de comidas al 10 % (Ley 37/1992, art. 91.Uno.2.2.º); anticipo del 30 %;
  · desplazamiento por km de ida y vuelta: línea aparte SIN margen; suplemento de show cooking sobre el precio;
  · tramos de cancelación de los contratos 1.2 según el canal (Particular sobre el total; Empresa sobre el anticipo);
  · margen bruto CALCULADO (nunca tecleado) sobre el subtotal de comida; 3 ejemplos por las mismas fórmulas; autotest pycel.
Hojas `Calculadora` · `Ejemplos` · `Instrucciones` (firma `calc-beta` de `motor.detectar`). A4, protegida sin
contraseña, verdes desbloqueadas, línea de versión y bio de la familia. Corre en el VPS (openpyxl + pycel).

    /root/venv-guias/bin/python calculadora_pae_es.py --salida <dir>      # genera + inject_cache + autotest
    /root/venv-guias/bin/python calculadora_pae_es.py --salida <dir> --autotest

Sesión Claude Code, 11-oct-2026.
"""
import argparse
import math
import os
import re
import sys
from collections import OrderedDict

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill, Protection
from openpyxl.worksheet.datavalidation import DataValidation

AQUI = os.path.dirname(os.path.abspath(__file__))
PD = os.path.dirname(AQUI)
sys.path.insert(0, os.path.join(PD, 'planes-v2_0'))
import motor                                                                       # noqa: E402

PID = 'plan-negocio-paellero-eventos'
FICHERO = 'calculadora-pricing-eventos-paellero.xlsx'
TITULO = 'Calculadora de precios de paella para eventos (del coste al presupuesto)'
H1, H2, H3 = 'Calculadora', 'Ejemplos', 'Instrucciones'
VERDE = PatternFill('solid', fgColor='E8F5E9')
GRIS = PatternFill('solid', fgColor='F2F2F2')
FUENTE = 'Calibri'
CANALES = ['Particular', 'Empresa']
ARROCES = ['Valenciana', 'Marisco', 'Senyoret', 'Arroz negro', 'Mixta', 'A banda', 'Fideuá', 'Meloso de bogavante',
           'Caldoso de bogavante', 'Montaña', 'Alicantina con costra', 'Vegetal']
EUR, ENT, PCT, DEC = '#,##0.00 €', '#,##0', '0%', '#,##0.0'
GESTION = 200.0                                                                    # gestión de cancelación (Empresa)

# diámetro (cm) -> (raciones, gramos de arroz, litros de caldo): manual técnico 1.2, tabla 1
TABLA = OrderedDict([(50, (7, 700, 1.8)), (60, (11, 1100, 2.8)), (70, (16, 1600, 4.0)), (80, (22, 2200, 5.5)),
                     (90, (30, 3000, 7.5)), (100, (38, 3800, 9.5)), (120, (60, 6000, 15.0)), (150, (90, 9000, 22.5))])

ENTRADAS = [
    ('canal', 'Cliente', '@', 'Particular',
     'Particular o Empresa: cambia los tramos de cancelación (contratos del kit). Ayuntamientos y comisiones de '
     'fiestas cuentan como Empresa.'),
    ('guests', 'Invitados previstos', ENT, 120, 'Número entero de raciones que sirves.'),
    ('arroz', 'Arroz principal (informativo)', '@', 'Valenciana',
     'Sólo etiqueta el presupuesto: no entra en el cálculo. Cada variedad extra suma su paellera.'),
    ('diam', 'Diámetro de paellera (cm)', ENT, 120,
     'Fija las raciones, el arroz y el caldo por paellera con la tabla del final de esta hoja.'),
    ('mp', 'Coste de materia prima por invitado', EUR, 8.90,
     'Tu escandallo. La carta del kit da el food cost por ración con acompañamientos (7,60-16,50 €) para bodas y '
     'empresas; una paella sola de fiesta popular ronda 2,00-2,50 €.'),
    ('menaje', 'Coste de menaje, pan y desechables por invitado', EUR, 0.40,
     'Estimación del kit: 0,30-0,45 €. Pon 0 si el cliente pone el menaje.'),
    ('gas_kg', 'Gas por paellera y servicio (kg)', DEC, 1.3, 'Estimación del kit: 1,3 kg por paellera.'),
    ('gas_p', 'Precio del gas (€ por kg)', EUR, 2.25,
     'Propano de 35 kg a 78,80 € en julio de 2026 (boletín del IDAE): unos 2,25 € por kg.'),
    ('help', 'Ayudantes (número)', ENT, 2,
     'El manual pide ayudante desde 4 paelleras u 80 invitados. Los paelleros salen del cálculo.'),
    ('hours', 'Horas por persona (carga, viaje, montaje, cocción, servicio y recogida)', DEC, 9,
     'Estimación del kit: pon las horas reales de tu equipo.'),
    ('wage', 'Coste bruto por hora del personal', EUR, 13.0,
     'Estimación del kit. Copia tu convenio provincial de hostelería.'),
    ('ss', 'Seguridad Social a cargo de la empresa', PCT, 0.33,
     '30,65 % de contingencias, desempleo, FOGASA, formación y MEI, más AT/EP (fijos discontinuos, art. 16 ET).'),
    ('rentals', 'Importe de alquileres repercutidos (carpa, mesas, sillas)', EUR, 400.0,
     'Lo que te cobra el alquilador para este evento. Si lo facturas aparte sin ser accesorio, ese importe va al 21 % '
     'de IVA.'),
    ('km', 'Kilómetros de ida (obrador → espacio)', ENT, 35, 'Se cobran ida y vuelta, en una línea aparte.'),
    ('per_km', 'Tarifa por kilómetro de desplazamiento', EUR, 0.95,
     'La del contrato B2B del kit: se cobra ida y vuelta y va sin margen.'),
    ('show', 'Suplemento de show cooking (% sobre el precio por invitado)', PCT, 0.0,
     'Un 8 % es el valor del kit si cocinas la paella a la vista de los invitados. 0 si no hay show.'),
    ('margin', 'Margen bruto objetivo sobre el subtotal', PCT, 0.40,
     'Parte del subtotal que queda después de los costes directos de arriba. Estimación del kit.'),
    ('iva', 'IVA del servicio de comidas', PCT, 0.10,
     'Ley 37/1992, art. 91.Uno.2.2.º: comida y bebida para consumir en el acto, alcohol incluido. Lo que facturas '
     'aparte sin ser accesorio (carpa, mobiliario) va al 21 %.'),
    ('deposit', 'Anticipo para bloquear la fecha (% del total)', PCT, 0.30,
     'El de los contratos del kit: 30 % al firmar. Particular: el resto el día del evento. Empresa: el resto a '
     '30 días de la factura.'),
]
K = OrderedDict((k, 5 + i) for i, (k, *_r) in enumerate(ENTRADAS))
F_CALC = 5 + len(ENTRADAS) + 1
SALIDAS = [
    ('rac', 'Raciones por paellera (tabla)', ENT, '{t_rac}', 'Busca el diámetro en la tabla del final.'),
    ('pan', 'Paelleras necesarias', ENT, 'roundup({guests}/{rac},0)', 'Invitados entre raciones por paellera, hacia arriba.'),
    ('pae', 'Paelleros necesarios', ENT, 'roundup({pan}/3,0)',
     'Tres paelleras son el techo de un paellero (manual técnico 1.2).'),
    ('kg_arroz', 'Arroz total (kg)', DEC, '{pan}*{t_arroz}/1000', 'Paelleras × arroz por paellera de la tabla.'),
    ('l_caldo', 'Caldo total (litros)', DEC, '{pan}*{t_caldo}', 'Paelleras × caldo por paellera de la tabla.'),
    ('kg_gas', 'Gas total (kg)', DEC, '{pan}*{gas_kg}', ''),
    ('c_mp', 'Coste de materia prima', EUR, '{mp}*{guests}', ''),
    ('c_menaje', 'Coste de menaje, pan y desechables', EUR, '{menaje}*{guests}', ''),
    ('c_gas', 'Coste del gas', EUR, '{kg_gas}*{gas_p}', ''),
    ('c_staff', 'Coste de personal (salario + Seguridad Social)', EUR, '({pae}+{help})*{hours}*{wage}*(1+{ss})',
     'Paelleros y ayudantes. Si el dueño cocina, se cuenta a precio de hora: es su coste de oportunidad.'),
    ('c_rent', 'Importe de alquileres', EUR, '{rentals}', ''),
    ('c_total', 'Coste directo total', EUR, '{c_mp}+{c_menaje}+{c_gas}+{c_staff}+{c_rent}', ''),
    ('c_guest', 'Coste directo por invitado', EUR, '{c_total}/{guests}', ''),
    ('price', 'PRECIO POR INVITADO SIN IVA', EUR, 'roundup({c_total}/(1-{margin})/{guests}*(1+{show})*2)/2',
     'Coste ÷ (1 − margen) ÷ invitados, más el show cooking, redondeado a los 0,50 € de arriba. Es el de las '
     'empresas («+ IVA»).'),
    ('price_iva', 'PRECIO POR INVITADO CON IVA', EUR, '{price}*(1+{iva})',
     'Es el que ve un particular: en una oferta a consumidores el precio va con impuestos incluidos.'),
    ('subtotal', 'Subtotal de comida sin IVA', EUR, '{price}*{guests}', 'Precio por invitado × invitados.'),
    ('desp', 'Desplazamiento (ida y vuelta, sin margen)', EUR, '{km}*2*{per_km}', 'Línea aparte del presupuesto.'),
    ('base', 'Base imponible', EUR, '{subtotal}+{desp}', ''),
    ('iva_amt', 'IVA', EUR, '{base}*{iva}', ''),
    ('total', 'TOTAL DEL PRESUPUESTO (con IVA)', EUR, '{base}+{iva_amt}', ''),
    ('dep_amt', 'Anticipo para bloquear la fecha', EUR, '{total}*{deposit}', ''),
    ('balance', 'Pendiente después del anticipo', EUR, '{total}-{dep_amt}',
     'Particular: el día del evento. Empresa: a 30 días de la factura.'),
    ('profit', 'Margen bruto en euros', EUR, '{subtotal}-{c_total}', 'Sobre el subtotal de comida, sin el desplazamiento.'),
    ('margin_calc', 'Margen bruto (calculado)', '0.0%', 'if({subtotal}==0,0,{profit}/{subtotal})',
     'Sale de los costes, nunca se teclea. Queda igual o por encima del objetivo por el redondeo del precio.'),
    ('can1', 'Importe retenido si cancela con más de 60 días', EUR, 'if({canal}=="Particular",{total}*0.3,0)',
     'Particular: el 30 % del total. Empresa: nada.'),
    ('can2', 'Importe retenido si cancela entre 60 y 30 días', EUR,
     'if({canal}=="Particular",{total}*0.5,{dep_amt}*0.3)',
     'Particular: el 50 % del total. Empresa: el 30 % del anticipo.'),
    ('can3', 'Importe retenido si cancela con menos de 30 días (Empresa: entre 30 y 15)', EUR,
     'if({canal}=="Particular",{total}*0.75,{dep_amt}*0.6+%d)' % GESTION,
     'Particular: el 75 % del total. Empresa: el 60 % del anticipo más 200 € de gestión.'),
    ('can4', 'Importe retenido si cancela con 72 horas o menos (Empresa: menos de 15 días)', EUR, '{total}',
     'Particular y Empresa: el 100 % del total.'),
]
S = OrderedDict((k, F_CALC + 1 + i) for i, (k, *_r) in enumerate(SALIDAS))
F_TABLA = F_CALC + 1 + len(SALIDAS) + 1

EJEMPLOS = OrderedDict([
    ('Boda en finca', dict(canal='Particular', guests=120, arroz='Valenciana', diam=120, mp=8.90, menaje=0.40,
                           gas_kg=1.3, gas_p=2.25, help=2, hours=9, wage=13.0, ss=0.33, rentals=400.0, km=35,
                           per_km=0.95, show=0.0, margin=0.40, iva=0.10, deposit=0.30)),
    ('Comida de empresa con show cooking', dict(canal='Empresa', guests=70, arroz='Marisco', diam=90, mp=8.50,
                                                menaje=0.45, gas_kg=1.3, gas_p=2.25, help=1, hours=7, wage=13.0,
                                                ss=0.33, rentals=0.0, km=25, per_km=0.95, show=0.08, margin=0.45,
                                                iva=0.10, deposit=0.30)),
    ('Fiesta popular de ayuntamiento', dict(canal='Empresa', guests=400, arroz='Mixta', diam=150, mp=2.25,
                                            menaje=0.30, gas_kg=1.3, gas_p=2.25, help=2, hours=7, wage=12.0, ss=0.33,
                                            rentals=0.0, km=30, per_km=0.95, show=0.0, margin=0.12, iva=0.10,
                                            deposit=0.30)),
])
EJ_NOTAS = OrderedDict([
    ('Boda en finca', 'Particular, 120 invitados, dos variedades en paellera de 120 cm más entrantes (la materia '
                      'prima incluye los acompañamientos), carpa y mesas alquiladas repercutidas y precio con IVA.'),
    ('Comida de empresa con show cooking', 'Empresa: presupuesto «+ IVA», cocina a la vista con el 8 % de suplemento '
                                           'y tramos de cancelación sobre el anticipo.'),
    ('Fiesta popular de ayuntamiento', 'Ayuntamiento = Empresa. 400 raciones en paelleras de 150 cm, menaje '
                                       'desechable y margen bajo por volumen: el precio por ración queda en el '
                                       'orden de 4-5 € (en 2025, Sayalonga cobró 1.480 € por 400 raciones).'),
])


def tabla_ref(col, ref_diam):
    rng_d = "'%s'!$A$%d:$A$%d" % (H1, F_TABLA + 1, F_TABLA + len(TABLA))
    rng = lambda c: "'%s'!$%s$%d:$%s$%d" % (H1, c, F_TABLA + 1, c, F_TABLA + len(TABLA))
    return 'INDEX(%s,MATCH(%s,%s,0))' % (rng(col), ref_diam, rng_d)


def refs(col):
    """Referencias a celdas de la columna `col` (B en Calculadora; B/C/D en Ejemplos) + búsquedas en la tabla."""
    ref = dict((k, '$%s$%d' % (col, f)) for k, f in K.items())
    ref.update((k, '$%s$%d' % (col, f)) for k, f in S.items())
    ref['t_rac'] = tabla_ref('B', ref['diam'])
    ref['t_arroz'] = tabla_ref('C', ref['diam'])
    ref['t_caldo'] = tabla_ref('D', ref['diam'])
    return ref


def calcular(e):
    r = OrderedDict()
    rac, g_arroz, l_caldo = TABLA[e['diam']]
    r['rac'] = rac
    r['pan'] = math.ceil(round(e['guests'] / rac, 9))
    r['pae'] = math.ceil(round(r['pan'] / 3, 9))
    r['kg_arroz'] = r['pan'] * g_arroz / 1000
    r['l_caldo'] = r['pan'] * l_caldo
    r['kg_gas'] = r['pan'] * e['gas_kg']
    r['c_mp'] = e['mp'] * e['guests']
    r['c_menaje'] = e['menaje'] * e['guests']
    r['c_gas'] = r['kg_gas'] * e['gas_p']
    r['c_staff'] = (r['pae'] + e['help']) * e['hours'] * e['wage'] * (1 + e['ss'])
    r['c_rent'] = e['rentals']
    r['c_total'] = r['c_mp'] + r['c_menaje'] + r['c_gas'] + r['c_staff'] + r['c_rent']
    r['c_guest'] = r['c_total'] / e['guests']
    r['price'] = math.ceil(round(r['c_total'] / (1 - e['margin']) / e['guests'] * (1 + e['show']) * 2, 9)) / 2
    r['price_iva'] = r['price'] * (1 + e['iva'])
    r['subtotal'] = r['price'] * e['guests']
    r['desp'] = e['km'] * 2 * e['per_km']
    r['base'] = r['subtotal'] + r['desp']
    r['iva_amt'] = r['base'] * e['iva']
    r['total'] = r['base'] + r['iva_amt']
    r['dep_amt'] = r['total'] * e['deposit']
    r['balance'] = r['total'] - r['dep_amt']
    r['profit'] = r['subtotal'] - r['c_total']
    r['margin_calc'] = 0 if r['subtotal'] == 0 else r['profit'] / r['subtotal']
    part = e['canal'] == 'Particular'
    r['can1'] = r['total'] * 0.3 if part else 0.0
    r['can2'] = r['total'] * 0.5 if part else r['dep_amt'] * 0.3
    r['can3'] = r['total'] * 0.75 if part else r['dep_amt'] * 0.6 + GESTION
    r['can4'] = r['total']
    return r


def a_excel(expr, ref):
    f = expr.replace('max(', 'MAX(').replace('roundup(', 'ROUNDUP(').replace('if(', 'IF(').replace('==', '=')
    f = re.sub(r'\{(\w+)\}', lambda m: ref[m.group(1)], f)
    return f.replace('*2)/2', '*2,0)/2')


def _f(ws, coord, v, b=False, sz=11, i=False, fmt=None, wrap=False, fill=None, locked=True):
    c = ws[coord]
    c.value = v
    c.font = Font(name=FUENTE, size=sz, bold=b, italic=i)
    if fmt:
        c.number_format = fmt
    c.alignment = Alignment(wrap_text=wrap, vertical='top')
    if fill:
        c.fill = fill
    c.protection = Protection(locked=locked)
    return c


def _imprimir(ws, horizontal=False):
    ws.page_setup.paperSize = 9                                                   # A4
    ws.page_setup.orientation = 'landscape' if horizontal else 'portrait'
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.oddFooter.center.text = 'aichef.pro/%s · página &P de &N' % PID
    ws.print_options.horizontalCentered = True
    ws.protection.sheet = True                                                    # sin contraseña
    ws.sheet_view.showGridLines = False


def _dv(ws, tipo, celdas, f1=None, f2=None, op=None, titulo='Valor no válido', msg=''):
    dv = DataValidation(type=tipo, formula1=f1, formula2=f2, operator=op, allow_blank=False, showErrorMessage=True,
                        errorTitle=titulo, error=msg)
    ws.add_data_validation(dv)
    for c in celdas:
        dv.add(c)


def hoja_calculadora(wb):
    ws = wb.active
    ws.title = H1
    for col, w in (('A', 62), ('B', 18), ('C', 80), ('D', 16)):
        ws.column_dimensions[col].width = w
    _f(ws, 'A1', 'CALCULADORA DE PRECIOS DE PAELLA PARA EVENTOS — DEL COSTE AL PRESUPUESTO', b=True, sz=14)
    _f(ws, 'A2', 'Escribe sólo en las celdas VERDES: todo lo demás se calcula. Importes sin IVA salvo donde la fila '
                 'diga lo contrario. Los valores por defecto son estimaciones del kit: cámbialos por los tuyos.',
       wrap=True)
    ws.merge_cells('A2:C2')
    ws.row_dimensions[2].height = 32
    for col, t in zip('ABC', ('DATOS DEL EVENTO', 'Valor', 'Notas')):
        _f(ws, '%s4' % col, t, b=True, fill=GRIS)
    for k, rot, fmt, val, nota in ENTRADAS:
        f = K[k]
        _f(ws, 'A%d' % f, rot)
        _f(ws, 'B%d' % f, val, fmt=fmt if fmt != '@' else None, fill=VERDE, locked=False)
        _f(ws, 'C%d' % f, nota, wrap=True)
    for col, t in zip('ABC', ('CÁLCULO', 'Valor', 'Notas')):
        _f(ws, '%s%d' % (col, F_CALC), t, b=True, fill=GRIS)
    ref = refs('B')
    for k, rot, fmt, expr, nota in SALIDAS:
        f = S[k]
        neg = k in ('price', 'price_iva', 'total', 'margin_calc')
        _f(ws, 'A%d' % f, rot, b=neg)
        _f(ws, 'B%d' % f, '=' + a_excel(expr, ref), fmt=fmt, b=neg)
        if nota:
            _f(ws, 'C%d' % f, nota, wrap=True)
    _f(ws, 'A%d' % F_TABLA, 'TABLA DE ESCALADO POR PAELLERA (manual técnico 1.2, tabla 1): diámetro en cm',
       b=True, fill=GRIS)
    for col, t in zip('BCD', ('Raciones', 'Arroz (g)', 'Caldo (L)')):
        _f(ws, '%s%d' % (col, F_TABLA), t, b=True, fill=GRIS)
    for i, (d, (rac, g, lc)) in enumerate(TABLA.items(), 1):
        _f(ws, 'A%d' % (F_TABLA + i), d, fmt='0" cm"')
        ws['A%d' % (F_TABLA + i)].alignment = Alignment(horizontal='left', vertical='top')
        _f(ws, 'B%d' % (F_TABLA + i), rac, fmt=ENT)
        _f(ws, 'C%d' % (F_TABLA + i), g, fmt=ENT)
        ws['C%d' % (F_TABLA + i)].alignment = Alignment(horizontal='left', vertical='top')
        _f(ws, 'D%d' % (F_TABLA + i), lc, fmt=DEC)
    f = F_TABLA + len(TABLA) + 1
    _f(ws, 'A%d' % f, 'Raciones = número bajo del rango del manual (ración generosa de evento). Un paellero atiende '
                      'como máximo tres paelleras.', i=True, wrap=True)
    ws.merge_cells('A%d:C%d' % (f, f))
    ws.row_dimensions[f].height = 32
    _dv(ws, 'list', [ref['canal'].replace('$', '')], f1='"%s"' % ','.join(CANALES), titulo='Cliente no válido',
        msg='Elige Particular o Empresa.')
    _dv(ws, 'list', [ref['arroz'].replace('$', '')], f1='"%s"' % ','.join(ARROCES), titulo='Arroz no válido',
        msg='Elige un arroz de la lista.')
    _dv(ws, 'list', [ref['diam'].replace('$', '')], f1='"%s"' % ','.join(str(d) for d in TABLA),
        titulo='Diámetro no válido', msg='Elige un diámetro de la tabla: 50, 60, 70, 80, 90, 100, 120 o 150 cm.')
    _dv(ws, 'whole', ['B%d' % K['guests']], f1='1', f2='100000', op='between', titulo='Número no válido',
        msg='Escribe un número entero de 1 o más.')
    _dv(ws, 'whole', ['B%d' % K[k] for k in ('help', 'km')], f1='0', f2='100000', op='between',
        titulo='Número no válido', msg='Escribe un número entero de 0 o más.')
    _dv(ws, 'decimal', ['B%d' % K[k] for k in ('mp', 'menaje', 'gas_kg', 'gas_p', 'hours', 'wage', 'rentals',
                                              'per_km')], f1='0', op='greaterThanOrEqual',
        titulo='Importe no válido', msg='Escribe un número mayor o igual que 0.')
    _dv(ws, 'decimal', ['B%d' % K[k] for k in ('ss', 'iva', 'deposit', 'show')], f1='0', f2='1', op='between',
        titulo='Porcentaje no válido', msg='Escribe un porcentaje entre 0 y 1 (0,10 = 10 %).')
    _dv(ws, 'decimal', ['B%d' % K['margin']], f1='0', f2='0.95', op='between', titulo='Porcentaje no válido',
        msg='Escribe un margen objetivo entre 0 y 95 % (0,40 = 40 %).')
    ws.freeze_panes = 'A5'
    _imprimir(ws)


def hoja_ejemplos(wb):
    ws = wb.create_sheet(H2)
    ws.column_dimensions['A'].width = 62
    for col in 'BCD':
        ws.column_dimensions[col].width = 26
    _f(ws, 'A1', '3 PRESUPUESTOS DE EJEMPLO — con las mismas fórmulas de la hoja «%s»' % H1, b=True, sz=14)
    _f(ws, 'A2', 'Estimaciones del kit para ilustrar el cálculo, no precios de mercado: cambia los datos en «%s» '
                 'para presupuestar tus eventos. La tabla de escalado por paellera está al final de esa hoja.'
       % H1, wrap=True)
    ws.merge_cells('A2:D2')
    ws.row_dimensions[2].height = 32
    _f(ws, 'A4', 'DATOS DEL EVENTO', b=True, fill=GRIS)
    for j, nombre in enumerate(EJEMPLOS):
        _f(ws, '%s4' % 'BCD'[j], nombre, b=True, fill=GRIS, wrap=True)
    ws.row_dimensions[4].height = 32
    for k, rot, fmt, val, nota in ENTRADAS:
        _f(ws, 'A%d' % K[k], rot)
        for j, e in enumerate(EJEMPLOS.values()):
            _f(ws, '%s%d' % ('BCD'[j], K[k]), e[k], fmt=fmt if fmt != '@' else None)
    _f(ws, 'A%d' % F_CALC, 'RESULTADO', b=True, fill=GRIS)
    for j in range(3):
        _f(ws, '%s%d' % ('BCD'[j], F_CALC), list(EJEMPLOS)[j], b=True, fill=GRIS, wrap=True)
    ws.row_dimensions[F_CALC].height = 32
    for k, rot, fmt, expr, nota in SALIDAS:
        neg = k in ('price', 'price_iva', 'total', 'margin_calc')
        _f(ws, 'A%d' % S[k], rot, b=neg)
        for j in range(3):
            col = 'BCD'[j]
            _f(ws, '%s%d' % (col, S[k]), '=' + a_excel(expr, refs(col)), fmt=fmt, b=neg)
    f = S[SALIDAS[-1][0]] + 2
    _f(ws, 'A%d' % f, 'QUÉ ENSEÑA CADA EJEMPLO', b=True, fill=GRIS)
    for j, (nombre, nota) in enumerate(EJ_NOTAS.items()):
        _f(ws, 'A%d' % (f + 1 + j), nombre, b=True)
        _f(ws, 'B%d' % (f + 1 + j), nota, wrap=True)
        ws.merge_cells('B%d:D%d' % (f + 1 + j, f + 1 + j))
        ws.row_dimensions[f + 1 + j].height = 48
    ws.freeze_panes = 'B5'
    _imprimir(ws, horizontal=True)


INSTRUCCIONES = [
    ('INSTRUCCIONES DE USO — Calculadora de precios de paella para eventos', True),
    ('Convierte tus costes en un precio por invitado y en un presupuesto completo: paelleras y paelleros, arroz, '
     'caldo y gas, subtotal, IVA, total, anticipo y lo que retienes si el cliente cancela. Sólo se escribe en las '
     'celdas VERDES; el margen se calcula.', False),
    ('1. Elige el cliente (Particular o Empresa; ayuntamientos y comisiones de fiestas cuentan como Empresa), los '
     'invitados y el diámetro de paellera.', False),
    ('2. La tabla de escalado del final de la primera hoja da las raciones, el arroz y el caldo por paellera. Las '
     'paelleras salen de dividir los invitados entre las raciones de la paellera, hacia arriba; un paellero atiende '
     'como máximo tres paelleras.', False),
    ('3. Escribe tu coste de materia prima por invitado (tu escandallo), el menaje, el gas y su precio.', False),
    ('4. Escribe los ayudantes, las horas por persona, el coste bruto por hora y la Seguridad Social. Si cocinas tú, '
     'cuéntate a precio de hora: es tu coste de oportunidad.', False),
    ('5. Escribe los alquileres que repercutes y los kilómetros de ida: el desplazamiento se cobra ida y vuelta, en '
     'una línea aparte y sin margen.', False),
    ('6. Fija el margen bruto objetivo y, si cocinas a la vista de los invitados, el suplemento de show cooking. El '
     'precio por invitado es el coste directo dividido entre (1 − margen), más el suplemento, redondeado a los '
     '0,50 € de arriba: el margen calculado queda igual o por encima del objetivo.', False),
    ('7. A una empresa se le presupuesta el precio SIN IVA («+ IVA»); a un particular, el precio CON IVA, porque en '
     'una oferta a consumidores el precio va con los impuestos incluidos.', False),
    ('IVA, PROPINAS Y COBROS', True),
    ('El servicio de comidas para consumir en el acto va al 10 % de IVA, alcohol incluido (Ley 37/1992, art. '
     '91.Uno.2.2.º). Lo que facturas aparte sin ser accesorio, como la carpa o el mobiliario, va al 21 %.', False),
    ('La propina voluntaria es decisión del cliente y es del equipo: no entra en el precio, en la base imponible ni '
     'en el total.', False),
    ('Los pagos en efectivo con particulares sólo valen hasta 1.000 € (art. 7 de la Ley 7/2012).', False),
    ('ANTICIPO Y CANCELACIÓN', True),
    ('El anticipo es del 30 % al firmar. Importes retenidos según los contratos del kit. Particular, sobre el total: '
     'más de 60 días, el 30 %; entre 60 y 30 días, el 50 %; menos de 30 días, el 75 %; con 72 horas o menos, el '
     '100 %. Empresa: más de 60 días, nada; entre 60 y 30 días, el 30 % del anticipo; entre 30 y 15 días, el 60 % '
     'del anticipo más 200 € de gestión; menos de 15 días, el 100 % del total. En un servicio de comida con fecha '
     'no hay derecho de desistimiento (art. 103.l del TRLGDCU): por eso la política tiene que estar firmada.', False),
    ('ESTIMACIONES DEL KIT', True),
    ('Los valores por defecto y los tres ejemplos son estimaciones para ilustrar el cálculo, no precios de mercado. '
     'Cámbialos por tu escandallo, tus salarios, tus alquileres y tus kilómetros antes de presupuestar.', False),
]


def hoja_instrucciones(wb):
    ws = wb.create_sheet(H3)
    ws.column_dimensions['A'].width = 110
    f = 1
    for t, b in INSTRUCCIONES:
        _f(ws, 'A%d' % f, t, b=b, sz=13 if f == 1 else 11, wrap=True)
        if not b:
            ws.row_dimensions[f].height = max(16, 15 * math.ceil(len(t) / 120.0))
        f += 2 if b and f > 1 else 1
    f += 1
    for t in ('Para editar una celda que no esté en verde: Revisar → Desproteger hoja (no tiene contraseña).',
              'Más plantillas y kits del catálogo en aichef.pro/productos-digitales', motor.BIO_LINE,
              motor.VERSION_FMT.format(pid=PID)):
        _f(ws, 'A%d' % f, t, wrap=True)
        f += 1
    _imprimir(ws)


def generar(destino):
    wb = openpyxl.Workbook()
    hoja_calculadora(wb)
    hoja_ejemplos(wb)
    hoja_instrucciones(wb)
    pr = wb.properties
    pr.title = '%s · v%s' % (TITULO, motor.VERSION)
    pr.subject = 'Plan de Negocio Paellero para Eventos'
    pr.description = 'aichef.pro/%s' % PID
    pr.keywords = 'plan negocio paellero eventos, AI Chef Pro'
    pr.creator = pr.lastModifiedBy = motor.CREATOR
    os.makedirs(os.path.dirname(destino), exist_ok=True)
    wb.save(destino)
    sys.path.insert(0, PD)
    import inject_cache as IC                                                      # noqa: E402
    nf, inj, fail = IC.inject(destino)
    print('%s: formulas=%d cache_inyectado=%d fallos_pycel=%d' % (os.path.basename(destino), nf, inj, fail))
    return fail


PROHIBIDOS = [r'\bfreelance', r'validad[oa]s?', r'Espa[ñn]a,? 2026', r'food ?truck', r'catering tem[aá]tico',
              r'estaciones tem[aá]ticas', r'cocina central']


def autotest(destino):
    from pycel import ExcelCompiler
    err, n = [], 0
    xl = ExcelCompiler(filename=destino)
    wv = openpyxl.load_workbook(destino, data_only=True)
    wf = openpyxl.load_workbook(destino)
    casos = [(H2, 'BCD'[j], e, nombre) for j, (nombre, e) in enumerate(EJEMPLOS.items())]
    casos.append((H1, 'B', OrderedDict((k, v) for k, _r, _f2, v, _n in ENTRADAS), 'hoja principal'))
    for hoja, col, e, nombre in casos:
        esperado = calcular(e)
        for k, v in esperado.items():
            c = '%s%d' % (col, S[k])
            n += 1
            for lab, x in (('pycel', xl.evaluate("'%s'!%s" % (hoja, c))), ('caché', wv[hoja][c].value)):
                if not isinstance(x, (int, float)) or abs(float(x) - v) > 0.005:
                    err.append('%s · %s %s (%s): %s %r ≠ esperado %r' % (nombre, hoja, c, k, lab, x, v))
        if esperado['margin_calc'] + 1e-9 < e['margin']:
            err.append('%s: margen calculado %.4f < objetivo %.4f' % (nombre, esperado['margin_calc'], e['margin']))
    # casos propios: cambian celdas de la hoja principal y se recalculan con pycel
    base = OrderedDict((k, v) for k, _r, _f2, v, _n in ENTRADAS)
    propios = [('propio 1: Empresa, 50 cm, 25 invitados', dict(base, canal='Empresa', diam=50, guests=25, help=0,
                                                              hours=6, mp=7.60, rentals=0.0, show=0.08, margin=0.5,
                                                              km=0)),
               ('propio 2: Particular, 150 cm, 91 invitados', dict(base, canal='Particular', diam=150, guests=91,
                                                                  help=1, mp=11.0, menaje=0.0, margin=0.35, km=60,
                                                                  deposit=0.25))]
    for nombre, e in propios:
        esperado = calcular(e)
        xl2 = ExcelCompiler(filename=destino)
        for k in S:
            xl2.evaluate("'%s'!B%d" % (H1, S[k]))
        for k, v in e.items():
            if k == 'arroz':                                                       # informativo: no entra en el cálculo
                continue
            xl2.set_value("'%s'!B%d" % (H1, K[k]), v)
        for k, v in esperado.items():
            n += 1
            x = xl2.evaluate("'%s'!B%d" % (H1, S[k]))
            if not isinstance(x, (int, float)) or abs(float(x) - v) > 0.005:
                err.append('%s: %s pycel %r ≠ esperado %r' % (nombre, k, x, v))
    e = dict(EJEMPLOS['Comida de empresa con show cooking'])
    if calcular(dict(e, canal='Particular'))['can2'] == calcular(dict(e, canal='Empresa'))['can2']:
        err.append('el selector Particular/Empresa no cambia la cancelación')
    if motor.detectar(wf, FICHERO)['molde'] != 'calc-beta':
        err.append('motor.detectar no la reconoce como calc-beta')
    gf = motor.gate_formatos(wf, FICHERO)
    if gf:
        err.append('gate_formatos: %d celdas, p. ej. %r' % (len(gf), gf[:3]))
    # validaciones de datos
    ws = wf[H1]
    dvs = dict((str(rng), dv) for dv in ws.data_validations.dataValidation for rng in dv.sqref.ranges)
    for k, esperado in (('canal', 'Particular,Empresa'), ('diam', '50,60,70,80,90,100,120,150')):
        c = 'B%d' % K[k]
        dv = dvs.get(c)
        if dv is None or dv.type != 'list' or esperado not in (dv.formula1 or ''):
            err.append('DV de lista incorrecta en %s' % c)
    dv = dvs.get('B%d' % K['arroz'])
    if dv is None or dv.type != 'list' or any(a not in dv.formula1 for a in ARROCES):
        err.append('DV de lista del arroz incorrecta')
    dv = dvs.get('B%d' % K['margin'])
    if dv is None or dv.formula2 != '0.95':
        err.append('DV del margen (máx. 95 %)')
    for k in ('guests', 'help', 'km'):
        dv = dvs.get('B%d' % K[k])
        if dv is None or dv.type != 'whole':
            err.append('DV entera ausente en %s' % k)
    textos = []
    for w in wf.worksheets:
        if w.page_setup.paperSize != 9 or not w.protection.sheet:
            err.append('%s: A4 y protección' % w.title)
        for row in w.iter_rows():
            for c in row:
                v = c.value
                if isinstance(v, str) and c.data_type != 'f':
                    textos.append((w.title, c.coordinate, v))
                if c.data_type == 'f' and wv[w.title][c.coordinate].value is None:
                    err.append('%s!%s: fórmula sin caché' % (w.title, c.coordinate))
                if c.fill.fgColor.rgb in ('00E8F5E9', 'FFE8F5E9') and c.protection.locked:
                    err.append('%s!%s: verde bloqueada' % (w.title, c.coordinate))
    todo = '\n'.join(t for _h, _c, t in textos)
    for rx in PROHIBIDOS:
        if re.search(rx, todo, re.I):
            err.append('prohibido /%s/' % rx)
    if re.search(r'service charge|sales tax|\$', todo, re.I):
        err.append('restos del EN (service charge / sales tax / $)')
    if motor.VERSION_FMT.format(pid=PID) not in todo or motor.BIO_LINE not in todo:
        err.append('falta la línea de versión o la bio')
    if wf.sheetnames != [H1, H2, H3]:
        err.append('pestañas %r' % wf.sheetnames)
    p = wf.properties
    if p.creator != 'AI Chef Pro' or p.description != 'aichef.pro/%s' % PID or \
            p.keywords != 'plan negocio paellero eventos, AI Chef Pro':
        err.append('docProps: creator=%r description=%r keywords=%r' % (p.creator, p.description, p.keywords))
    for x in err:
        print('ERROR', x)
    print('autotest calculadora paellero ES: %d comprobaciones (pycel + caché en ejemplos y hoja principal, pycel en '
          '2 casos propios) · %s' % (n * 2, 'VERDE' if not err else 'ROJO (%d errores)' % len(err)))
    for nombre, e in EJEMPLOS.items():
        r = calcular(e)
        print('  %-36s %2d paelleras · %d paelleros · %6.2f €/inv sin IVA · %6.2f con IVA · total %9.2f · '
              'anticipo %8.2f · margen %.1f %%' % (nombre, r['pan'], r['pae'], r['price'], r['price_iva'], r['total'],
                                                   r['dep_amt'], r['margin_calc'] * 100))
    return 1 if err else 0


if __name__ == '__main__':
    ap = argparse.ArgumentParser()
    ap.add_argument('--salida', required=True, help='carpeta de salida')
    ap.add_argument('--autotest', action='store_true')
    a = ap.parse_args()
    destino = os.path.join(a.salida, FICHERO)
    if not a.autotest and generar(destino):
        sys.exit('ERROR: fallos de pycel en inject_cache')
    sys.exit(autotest(destino))
