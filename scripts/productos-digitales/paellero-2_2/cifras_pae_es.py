#!/usr/bin/env python3
"""
cifras_pae_es.py — Paellero ES 2.2 (copia de `catering-es-2_0/cifras_cat_es.py`): lee la CACHÉ del plan financiero (después de `inject_cache`, que hace `main.py`)
por RÓTULO, escribe `cifras_caso_es.json` (los tokens {ASI} del docx, D8) y pasa el gate de calibración D15:
DSCR mínimo ≥ 1,25 · saldo mínimo de tesorería > 0 · cobertura de horas ≥ 98 % · holgura sobre el equilibrio de caja
≥ 15 % · diferencia origen-usos ≥ 0 y < 5 %. Corre en el VPS (openpyxl).

    /root/venv-guias/bin/python cifras_pae_es.py --xlsx <plan-financiero-paellero-eventos.xlsx>

Sesión Claude Code, 11-oct-2026.
"""
import argparse
import json
import os
import re
import sys

import openpyxl

AQUI = os.path.dirname(os.path.abspath(__file__))
SALIDA = os.path.join(AQUI, 'cifras_caso_es.json')

#: token → (hoja, regex del rótulo de la columna A, columna, formato). Formatos: eur (miles con punto, sin
#: decimales) · eur2 · pct (1 decimal) · pct0 · dec1 · dec2 · int.
MAPA = {
    'INVITADOS_EVENTO': ('0. Supuestos', r'^invitados por evento \(media', 2, 'int'),
    'PRECIO_INVITADO': ('0. Supuestos', r'^precio medio por invitado sin iva', 2, 'eur'),
    'PRECIO_INVITADO_IVA': ('0. Supuestos', r'^pvp equivalente con iva', 2, 'eur2'),
    'EVENTOS_ANO': ('0. Supuestos', r'^eventos al a[ñn]o$', 2, 'int'),
    'CREC_A2': ('0. Supuestos', r'^crecimiento de volumen del a[ñn]o 2', 2, 'pct0'),
    'CREC_A3': ('0. Supuestos', r'^crecimiento de volumen del a[ñn]o 3', 2, 'pct0'),
    'PCT_COMIDA': ('0. Supuestos', r'^ventas de comida sobre el total', 2, 'pct0'),
    'FOOD_COST': ('0. Supuestos', r'^coste de mercanc[ií]a sobre las ventas de comida', 2, 'pct0'),
    'BEBIDA_COST': ('0. Supuestos', r'^coste de mercanc[ií]a sobre las ventas de bebida', 2, 'pct0'),
    'CUOTA_COCINA_MES': ('0. Supuestos', r'^cuota mensual de la cocina de apoyo', 2, 'eur'),
    'SEGUROS_ANO': ('0. Supuestos', r'^seguros', 2, 'eur'),
    'RECURSOS_PROPIOS': ('0. Supuestos', r'^recursos propios', 2, 'eur'),
    'TIPO_PRESTAMO': ('0. Supuestos', r'^tipo de inter[eé]s nominal', 2, 'pct'),
    'PLAZO_PRESTAMO': ('0. Supuestos', r'^plazo del pr[eé]stamo', 2, 'int'),
    'CARENCIA': ('0. Supuestos', r'^carencia de principal', 2, 'int'),
    'DIAS_COBRO': ('0. Supuestos', r'^d[ií]as medios de cobro', 2, 'int'),
    'CAPACIDAD_EVENTO': ('0. Supuestos', r'^capacidad m[aá]xima por evento', 2, 'int'),
    'VENTAS_A1': ('PyG 3 Años', r'^ingresos totales', 2, 'eur'),
    'VENTAS_A2': ('PyG 3 Años', r'^ingresos totales', 3, 'eur'),
    'VENTAS_A3': ('PyG 3 Años', r'^ingresos totales', 4, 'eur'),
    'COSTES_VARIABLES_A1': ('PyG 3 Años', r'^total costes variables', 2, 'eur'),
    'MARGEN_BRUTO_A1': ('PyG 3 Años', r'^margen bruto$', 2, 'eur'),
    'MARGEN_BRUTO_PCT_A1': ('PyG 3 Años', r'^margen bruto / ventas', 2, 'pct'),
    'COCINA_ANO': ('PyG 3 Años', r'^cuota de la cocina de apoyo', 2, 'eur'),
    'PERSONAL_A1': ('PyG 3 Años', r'^personal \(n[oó]minas', 2, 'eur'),
    'PERSONAL_PCT_A1': ('PyG 3 Años', r'^coste de personal / ventas', 2, 'pct'),
    'MERCANCIA_PCT_A1': ('PyG 3 Años', r'^coste de mercanc[ií]a / ventas', 2, 'pct'),
    'MARKETING_ANO': ('PyG 3 Años', r'^marketing y publicidad', 2, 'eur'),
    'AMORTIZACION_A1': ('PyG 3 Años', r'^amortizaci[oó]n del inmovilizado', 2, 'eur'),
    'INTERESES_A1': ('PyG 3 Años', r'^gastos financieros', 2, 'eur'),
    'FIJOS_A1': ('PyG 3 Años', r'^total costes fijos', 2, 'eur'),
    'RAI_A1': ('PyG 3 Años', r'^resultado antes de impuestos', 2, 'eur'),
    'RAI_A2': ('PyG 3 Años', r'^resultado antes de impuestos', 3, 'eur'),
    'RAI_A3': ('PyG 3 Años', r'^resultado antes de impuestos', 4, 'eur'),
    'NETO_A1': ('PyG 3 Años', r'^resultado neto$', 2, 'eur'),
    'NETO_A2': ('PyG 3 Años', r'^resultado neto$', 3, 'eur'),
    'NETO_A3': ('PyG 3 Años', r'^resultado neto$', 4, 'eur'),
    'NETO_PCT_A1': ('PyG 3 Años', r'^resultado neto / ventas', 2, 'pct'),
    'NETO_PCT_A3': ('PyG 3 Años', r'^resultado neto / ventas', 4, 'pct'),
    'CAPEX': ('Inversión Inicial', r'^capex', 2, 'eur'),
    'FONDO_MANIOBRA': ('Inversión Inicial', r'^fondo de maniobra$', 2, 'eur'),
    'INVERSION_TOTAL': ('Inversión Inicial', r'^inversi[oó]n total', 2, 'eur'),
    'IVA_INVERSION': ('Inversión Inicial', r'^iva soportado sobre la inversi', 2, 'eur'),
    'NECESIDAD_CAJA': ('Inversión Inicial', r'^necesidad total de caja', 2, 'eur'),
    'FURGONETA': ('Inversión Inicial', r'^furgoneta isoterma', 2, 'eur'),
    'PAELLERAS': ('Inversión Inicial', r'^paelleras de 60', 2, 'eur'),
    'QUEMADORES': ('Inversión Inicial', r'^quemadores de anillos', 2, 'eur'),
    'PRESTAMO': ('Financiación', r'^pr[eé]stamo bancario$', 2, 'eur'),
    'DIFERENCIA_FONDOS': ('Financiación', r'^diferencia \(origen - usos\)', 2, 'eur'),
    'DIFERENCIA_FONDOS_PCT': ('Financiación', r'^diferencia sobre los usos', 2, 'pct'),
    'CUOTA_ANUAL': ('Financiación', r'^cuota anual durante la amortizaci', 2, 'eur'),
    'DSCR_A2': ('Financiación', r'^dscr del a[ñn]o 2', 2, 'dec2'),
    'DSCR_MIN': ('Financiación', r'^dscr m[ií]nimo de todo el cuadro', 2, 'dec2'),
    'SALDO_MINIMO': ('Tesorería 12 meses', r'^saldo m[ií]nimo del a[ñn]o', 2, 'eur'),
    'MES_SALDO_MINIMO': ('Tesorería 12 meses', r'^mes en el que la caja toca fondo', 2, 'int'),
    'PAYBACK': ('Tesorería 12 meses', r'^payback del proyecto', 2, 'dec1'),
    'PAYBACK_CAPEX': ('Tesorería 12 meses', r'^payback sobre el capex', 2, 'dec1'),
    'BE_INVITADOS_EVENTO': ('Punto Equilibrio', r'^invitados necesarios por evento$', 2, 'int_up'),
    'BE_INVITADOS_EVENTO_CAJA': ('Punto Equilibrio', r'^invitados necesarios por evento \(caja\)', 2, 'int_up'),
    'BE_INGRESOS': ('Punto Equilibrio', r'^ingresos necesarios al a[ñn]o$', 2, 'eur'),
    'BE_INGRESOS_CAJA': ('Punto Equilibrio', r'^ingresos necesarios al a[ñn]o \(caja\)', 2, 'eur'),
    'HOLGURA_CAJA': ('Punto Equilibrio', r'^holgura sobre el equilibrio de caja', 2, 'pct'),
    'HOLGURA_CAJA_A2': ('Punto Equilibrio', r'^holgura sobre el equilibrio de caja', 3, 'pct'),
    'HOLGURA_CAJA_A3': ('Punto Equilibrio', r'^holgura sobre el equilibrio de caja', 4, 'pct'),
    'BE_INGRESOS_CAJA_A2': ('Punto Equilibrio', r'^ingresos necesarios al a[ñn]o \(caja\)', 3, 'eur'),
    'INVITADOS_EVENTO_A2': ('PyG 3 Años', r'^invitados por evento \(media\)', 3, 'dec1'),
    'INVITADOS_EVENTO_A3': ('PyG 3 Años', r'^invitados por evento \(media\)', 4, 'dec1'),
    'COBERTURA_HORAS': ('Personal', r'^cobertura \(contratadas', 2, 'pct0'),
    'PLANTILLA_PERSONAS': ('Personal', r'^total plantilla', 2, 'int'),
    'ESC_PES_VENTAS': ('Escenarios', r'^ingresos anuales', 2, 'eur'),
    'ESC_OPT_VENTAS': ('Escenarios', r'^ingresos anuales', 4, 'eur'),
    'ESC_PES_NETO': ('Escenarios', r'^resultado neto$', 2, 'eur'),
    'ESC_OPT_NETO': ('Escenarios', r'^resultado neto$', 4, 'eur'),
    'ESC_PES_EVENTOS': ('Escenarios', r'^eventos al a[ñn]o', 2, 'int'),
    'ESC_OPT_EVENTOS': ('Escenarios', r'^eventos al a[ñn]o', 4, 'int'),
}


def fmt(v, f):
    if f in ('int', 'int_up'):
        import math
        n = math.ceil(v - 1e-9) if f == 'int_up' else round(v)
        return '{:,}'.format(int(n)).replace(',', '.')
    if f == 'eur':
        return '{:,}'.format(int(round(v))).replace(',', '.') + ' €'
    if f == 'eur2':
        return '{:,.2f}'.format(v).replace(',', 'X').replace('.', ',').replace('X', '.') + ' €'
    if f in ('pct', 'pct0'):
        txt = ('{:.1f}' if f == 'pct' else '{:.0f}').format(v * 100).replace('.', ',')
        return txt + ' %'
    if f in ('dec1', 'dec2'):
        return ('{:.1f}' if f == 'dec1' else '{:.2f}').format(v).replace('.', ',')
    raise ValueError(f)


def leer(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    crudo, falta, textos = {}, [], {}
    for tok, (hoja, rx, col, f) in MAPA.items():
        ws = wb[hoja]
        v = None
        for r in range(1, ws.max_row + 1):
            a = ws.cell(r, 1).value
            if isinstance(a, str) and re.search(rx, a.strip(), re.I):
                v = ws.cell(r, col).value
                break
        if isinstance(v, str) and tok.startswith('PAYBACK') and v.strip():   # «Más de 3 años»: texto del libro
            textos[tok] = v.strip()
            continue
        if not isinstance(v, (int, float)):
            falta.append('%s (%s · %s): %r' % (tok, hoja, rx, v))
            continue
        crudo[tok] = (float(v), f)
    # derivados (D2/D8): los eventos al año de equilibrio con el tamaño medio de evento del plan
    if all(k in crudo for k in ('BE_INGRESOS', 'BE_INGRESOS_CAJA', 'PRECIO_INVITADO', 'INVITADOS_EVENTO')):
        ticket_evento = crudo['PRECIO_INVITADO'][0] * crudo['INVITADOS_EVENTO'][0]
        crudo['TICKET_EVENTO'] = (ticket_evento, 'eur')
        crudo['BE_EVENTOS_ANO'] = (crudo['BE_INGRESOS'][0] / ticket_evento, 'int_up')
        crudo['BE_EVENTOS_ANO_CAJA'] = (crudo['BE_INGRESOS_CAJA'][0] / ticket_evento, 'int_up')
        crudo['BE_EVENTOS_MES'] = (crudo['BE_INGRESOS'][0] / ticket_evento / 12, 'dec1')
        if 'BE_INGRESOS_CAJA_A2' in crudo:   # REVISION-FINAL M3: equilibrio de caja del año 2 sin crecer
            crudo['BE_EVENTOS_CAJA_A2'] = (crudo['BE_INGRESOS_CAJA_A2'][0] / ticket_evento, 'int_up')
    return crudo, falta, textos


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--xlsx', required=True)
    ap.add_argument('--salida', default=SALIDA)
    a = ap.parse_args()
    crudo, falta, textos = leer(a.xlsx)
    tokens = dict((k, fmt(v, f)) for k, (v, f) in crudo.items())
    tokens.update(textos)
    json.dump({'fuente': os.path.basename(a.xlsx), 'tokens': tokens,
               'valores': dict((k, v) for k, (v, _f) in crudo.items())},
              open(a.salida, 'w', encoding='utf-8'), ensure_ascii=False, indent=1, sort_keys=True)
    v = dict((k, x[0]) for k, x in crudo.items())
    gates = [
        ('DSCR mínimo ≥ 1,25', v.get('DSCR_MIN', 0) >= 1.25),
        ('saldo mínimo de tesorería > 0', v.get('SALDO_MINIMO', -1) > 0),
        ('cobertura de horas ≥ 98 %', v.get('COBERTURA_HORAS', 0) >= 0.98),
        ('holgura sobre el equilibrio de caja ≥ 15 %', v.get('HOLGURA_CAJA', 0) >= 0.15),
        ('holgura de caja del año 2 (sin carencia) ≥ 0 %', v.get('HOLGURA_CAJA_A2', -1) >= 0),
        ('origen - usos ≥ 0 y < 5 %', 0 <= v.get('DIFERENCIA_FONDOS_PCT', -1) < 0.05),
    ]
    for tok in ('EVENTOS_ANO', 'INVITADOS_EVENTO', 'PRECIO_INVITADO', 'VENTAS_A1', 'RAI_A1', 'NETO_A1',
                'INVERSION_TOTAL', 'NECESIDAD_CAJA', 'PRESTAMO', 'DSCR_MIN', 'SALDO_MINIMO', 'BE_INVITADOS_EVENTO',
                'BE_INVITADOS_EVENTO_CAJA', 'BE_EVENTOS_ANO', 'BE_EVENTOS_ANO_CAJA', 'HOLGURA_CAJA',
                'COBERTURA_HORAS', 'PAYBACK', 'HOLGURA_CAJA_A2', 'BE_EVENTOS_CAJA_A2', 'DSCR_A2'):
        print('  %-26s %s' % (tok, tokens.get(tok, '—')))
    for nombre, ok in gates:
        print('  [%s] %s' % ('OK' if ok else 'ROJO', nombre))
    for f in falta:
        print('  FALTA', f)
    print('cifras_caso_es.json: %d tokens · calibración D15 %s' % (
        len(tokens), 'VERDE' if all(ok for _n, ok in gates) and not falta else 'ROJA'))
    sys.exit(0 if all(ok for _n, ok in gates) and not falta else 1)


if __name__ == '__main__':
    main()
