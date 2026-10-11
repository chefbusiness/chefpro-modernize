#!/usr/bin/env python3
"""barrido_es.py — Catering ES 2.0 (SPEC D12): barre los xlsx/docx que se le pasen buscando vocabulario del molde FT,
del motor sin traducir y de la 1.x (freelance, validado, España 2026, cifras prohibidas). Lista celda/párrafo y texto.
Corre en el VPS. Uso: barrido_es.py [--prohibidos] <fichero>… · exit 1 si hay coincidencias. Sesión Claude Code, 11-oct-2026."""
import re
import sys

import docx
import openpyxl

sys.path.insert(0, __import__('os').path.dirname(__import__('os').path.abspath(__file__)))
import tandas_docx as T                                                            # noqa: E402

RX = re.compile(r'food ?truck|veh[ií]cul|cami[oó]n|ventanilla|generador|aparcamiento|\bcalle\b|'
                r'\bclientes?/d[ií]a|\bal d[ií]a\b|\bd[ií]as de apertura|\bdel local\b|\bel local\b|\ben sala\b|'
                r'\bticket\b|\bdelivery\b|\bcubiert|\bcomensal|' + '|'.join(T.PROHIBIDOS), re.I)
if '--prohibidos' in sys.argv:          # sólo lo prohibido por la SPEC (D12), para los 7 ficheros de texto libre
    RX = re.compile('|'.join(T.PROHIBIDOS), re.I)
n = 0
for p in [x for x in sys.argv[1:] if not x.startswith('--')]:
    if p.endswith('.xlsx'):
        wb = openpyxl.load_workbook(p)
        it = ((ws.title + '!' + c.coordinate, c.value) for ws in wb.worksheets for row in ws.iter_rows() for c in row
              if isinstance(c.value, str))
    else:
        d = docx.Document(p)
        it = (('p%d' % i, x.text) for i, x in enumerate(d.paragraphs))
    for loc, t in it:
        txt = ' '.join(re.findall(r'"([^"]*)"', t)) if t.startswith('=') else t
        m = RX.search(txt)
        if m:
            n += 1
            print('%s · %s · «%s» · …%s…' % (p.split('/')[-1][:28], loc, m.group(0), txt[max(0, m.start() - 50):m.end() + 50]))
print('barrido: %d coincidencias' % n)
sys.exit(1 if n else 0)
