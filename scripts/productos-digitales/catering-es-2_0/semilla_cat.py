#!/usr/bin/env python3
"""
semilla_cat.py — Catering ES 2.0 (SPEC D1/D3/D7): prepara el ORIGEN que `planes-v2_0/main.py` convierte en el plan
financiero y el checklist del catering.

Copia los dos xlsx del Food Truck ES v2.2 publicados (`dl/plan-negocio-food-truck/`, SOLO LECTURA) con los nombres de
fichero del catering y escribe lo que el módulo de contenido no alcanza (F1-camino §1):
  · plan: títulos de hoja heredados del fichero (Inversión A1/A2, PyG A2, Punto Equilibrio A2), las 14 partidas de
    inversión (filas 7-20: rótulo, importe, nota, ¿IVA?) y los costes fijos propios del P&L (filas 26-35);
  · checklist: título, pestañas, cabeceras de fase y las 82 tareas de `datos_checklist.py` (12/15/14/10/16/15) con el
    pie «Tareas completadas:» debajo; DV, formato condicional y contadores los rehace el motor (`checklist_ok_y_contador`).
Corre en el VPS (openpyxl). Después: `main.py --producto plan-catering-tematico-eventos --dry-run --origen <salida>`.

    /root/venv-guias/bin/python semilla_cat.py --salida /root/tmp-cates-f2/origen

Sesión Claude Code, 11-oct-2026.
"""
import argparse
import copy
import os
import shutil

import importlib
import sys

import openpyxl

import datos_checklist as DC

AQUI = os.path.dirname(os.path.abspath(__file__))
REPO = os.path.abspath(os.path.join(AQUI, '..', '..', '..'))
FT = os.path.join(REPO, 'astro-site', 'public', 'dl', 'plan-negocio-food-truck')
PLAN_FT, CHECK_FT = 'plan-financiero-food-truck.xlsx', 'checklist-apertura-food-truck.xlsx'
PLAN, CHECK = 'plan-financiero-catering-tematico-eventos.xlsx', 'checklist-apertura-catering-tematico-eventos.xlsx'

TITULOS_PLAN = {
    ('Inversión Inicial', 'A1'): 'PLAN DE NEGOCIO: CATERING TEMÁTICO DE EVENTOS',
    ('Inversión Inicial', 'A2'): 'Inversión inicial estimada — catering con estaciones temáticas en directo',
    ('PyG 3 Años', 'A2'): 'Escenario realista — catering temático de eventos',
    ('Punto Equilibrio', 'A2'): 'Punto de equilibrio — catering temático de eventos',
}

#: Filas 7-20 de 'Inversión Inicial' (molde FT): (rótulo, importe, nota, ¿lleva IVA?). Research §3: 67.000 € en 15
#: partidas; «web, fotos, marca y dossier» (4.000) y «lanzamiento» (3.000) van juntas para caber en las 14 filas.
#: Sin palabras de recuento en el rótulo (motor.RX_RECUENTO).
INVERSION = [
    ('Furgoneta frigorífica usada con certificado ATP', 22000,
     'Furgoneta de 2021-2022 con equipo de frío y certificado ATP vigente (RD 237/2000). Anuncios de segunda mano '
     'de octubre de 2026: 18.000-22.000 €', 'Sí'),
    ('Rotulación y estanterías de la furgoneta', 1500, 'Rotulación con la marca y estanterías con amarres', 'Sí'),
    ('Contenedores isotérmicos GN y carros de transporte', 3500,
     'Mantienen la cadena de frío y de caliente desde la cocina hasta la estación', 'Sí'),
    ('Equipo de producción: gastronorms, roner, sondas, básculas y cuchillería', 4000,
     'Lo que no aporta la cocina central compartida', 'Sí'),
    ('Estaciones temáticas portátiles (inducción, plancha y wok)', 9000,
     'Tres módulos de show cooking. El equipo específico de cada cocina del mundo se compra cuando esa estación '
     'se venda', 'Sí'),
    ('Chafing dishes y expositores de bufé', 3500, 'Para bufés, desayunos y coffee breaks', 'Sí'),
    ('Menaje base: vajilla, cubertería y cristalería', 6000,
     'Para 150 invitados; lo que pase se alquila y se repercute en el presupuesto', 'Sí'),
    ('Mantelería y uniformes', 2000, 'Con recambios para dos eventos seguidos', 'Sí'),
    ('Mobiliario de catering: mesas altas y barra plegable', 2500,
     'El resto del mobiliario y las carpas se alquilan para cada evento', 'Sí'),
    ('Marca, web, fotos, dossier de menús y lanzamiento', 7000,
     'Identidad, web con formulario de presupuesto, sesión de fotos y vídeo, dossier, ferias de novios, '
     'directorios y degustaciones de lanzamiento', 'Sí'),
    ('Software, tablet y datáfono', 1200, 'Presupuestos, agenda y facturación conforme a Veri*factu', 'Sí'),
    ('Constitución SL + notaría + registro', 1000,
     'SL unipersonal. Las tasas y el arancel del Registro no llevan IVA; los honorarios del notario y de la '
     'gestoría sí: cambia la columna de al lado a «Sí» si tu presupuesto es sobre todo de honorarios', 'No'),
    ('Declaración responsable, plan APPCC, alérgenos y formación', 1800,
     'Redacción del plan APPCC, fichas de alérgenos, formación del equipo y trámites sanitarios', 'Sí'),
    ('Stock inicial de producto', 2000, 'Producto seco, bebida y desechables para los primeros eventos', 'Sí'),
]
assert len(INVERSION) == 14 and sum(x[1] for x in INVERSION) == 67000

#: Filas 26-35 de 'PyG 3 Años' (costes fijos propios). Research §3: 22.150 €/año en 8 partidas; las dos filas
#: restantes del molde se llenan con dos costes reales del catering [estimado] en vez de suprimirlas: suprimir filas
#: desplaza el P&L y el motor lo cuenta como fórmulas perdidas (medido en el 2.º dry-run, 18 fórmulas). Sin «eventos»,
#: «clientes» ni «días» en el rótulo (motor.RX_RECUENTO).
FIJOS = {
    26: ('Asesoría fiscal y laboral (nóminas de fijos discontinuos)', 3000),
    27: ('Furgoneta: mantenimiento, ITV e inspección ATP', 1500),
    28: ('Marketing y publicidad', 9000),
    29: ('Software de gestión comercial, CRM y teléfono', 1800),
    30: ('Lavandería de mantelería y uniformes', 3000),
    31: ('Pruebas de menú y degustaciones comerciales', 3000),
    32: ('Gestión de residuos (orgánico, cartón y aceite usado)', 400),
    33: ('Reposición de menaje y roturas', 1200),
    34: ('Prevención de riesgos laborales y vigilancia de la salud', 450),
    35: ('Permisos y autorizaciones municipales', 600),
}


def _supuestos(wb):
    """Los valores de `0. Supuestos` del caso catering. `grupo_a` CONSERVA lo que el libro ya trae en las celdas
    verdes (§1.3: dato del cliente) y sólo usa el valor del contenido en una celda vacía: partiendo del libro del
    FT, sin este paso el caso seguiría siendo el del food truck (45 × 12 € × 250, medido en el primer dry-run)."""
    p2 = os.path.join(os.path.dirname(AQUI), 'planes-v2_0')
    sys.path.insert(0, p2)
    import grupo_a
    contenido = importlib.import_module('contenido_plan_catering_tematico_eventos.a')
    coords = dict((t[0], t[1]) for t in grupo_a.SUPUESTOS_BASE)
    ws = wb['0. Supuestos']
    escritas = 0
    for clave, t in contenido.SUPUESTOS.items():
        if t[2] is None:
            continue
        assert clave in coords, clave
        ws[coords[clave]] = t[2]
        escritas += 1
    return escritas


def plan(salida):
    destino = os.path.join(salida, PLAN)
    shutil.copy2(os.path.join(FT, PLAN_FT), destino)
    wb = openpyxl.load_workbook(destino)
    print('  supuestos escritos:', _supuestos(wb))
    for (hoja, coord), texto in TITULOS_PLAN.items():
        wb[hoja][coord] = texto
    ws = wb['Inversión Inicial']
    for i, (rot, imp, nota, iva) in enumerate(INVERSION):
        r = 7 + i
        assert not str(ws.cell(r, 1).value).lower().startswith(('fianza', 'alquiler', 'imprevistos')), r
        ws.cell(r, 1).value, ws.cell(r, 2).value, ws.cell(r, 4).value, ws.cell(r, 5).value = rot, imp, nota, iva
    ws = wb['PyG 3 Años']
    assert str(ws['A25'].value).startswith('Seguros') and str(ws['A36'].value).startswith('Amortización')
    for r, (rot, imp) in FIJOS.items():
        ws.cell(r, 1).value, ws.cell(r, 2).value = rot, imp
        ws.cell(r, 6).value = None          # la nota la pone el contenido (FIJOS)
    wb.save(destino)
    return destino


def _copiar_estilo(origen, destino):
    destino.font, destino.border = copy.copy(origen.font), copy.copy(origen.border)
    destino.fill, destino.alignment = copy.copy(origen.fill), copy.copy(origen.alignment)
    destino.number_format, destino.protection = origen.number_format, copy.copy(origen.protection)


def checklist(salida):
    destino = os.path.join(salida, CHECK)
    shutil.copy2(os.path.join(FT, CHECK_FT), destino)
    wb = openpyxl.load_workbook(destino)
    fases = [ws for ws in wb.worksheets if ws.title[:2] in ('F1', 'F2', 'F3', 'F4', 'F5', 'F6')]
    assert len(fases) == 6
    for ws, (titulo, cabecera, tareas) in zip(fases, DC.FASES):
        ws.title = titulo
        ws['A1'], ws['A2'] = DC.TITULO, cabecera
        assert ws['B4'].value and 'TRÁMITE' in ws['B4'].value.upper(), ws['B4'].value
        modelo = [copy.copy(ws.cell(5, c)) for c in range(1, 6)]
        alto = ws.row_dimensions[5].height
        pie = [copy.copy(ws.cell(r, c)) for r in range(5, ws.max_row + 1) for c in range(1, 6)
               if isinstance(ws.cell(r, c).value, str) and 'completad' in ws.cell(r, c).value.lower()]
        assert pie, ws.title
        fila_pie = pie[0].row
        estilo_pie = [copy.copy(ws.cell(fila_pie, c)) for c in range(1, 6)]
        # vaciar la zona de tareas y el pie (valores, DV y formato condicional los rehace el motor)
        for r in range(5, ws.max_row + 1):
            for c in range(1, 6):
                ws.cell(r, c).value = None
            ws.row_dimensions[r].height = None
        ws.data_validations.dataValidation = []
        ws.conditional_formatting = type(ws.conditional_formatting)()
        for i, (tarea, resp, plazo, nota) in enumerate(tareas):
            r = 5 + i
            for c, v in enumerate(('☐', tarea, resp, plazo, nota), start=1):
                cel = ws.cell(r, c)
                _copiar_estilo(modelo[c - 1], cel)
                cel.value = v
            ws.row_dimensions[r].height = alto
        r = 5 + len(tareas) + 1
        for c in range(1, 6):
            _copiar_estilo(estilo_pie[c - 1], ws.cell(r, c))
        ws.cell(r, 2).value = 'Tareas completadas:'
        ws.cell(r, 3).value = '=COUNTIF(A5:A%d,"✓")' % (4 + len(tareas))
        ws.cell(r, 4).value = 'de'
        ws.cell(r, 5).value = '=COUNTIF(B5:B%d,"?*")-COUNTIF(A5:A%d,"N/A")' % (4 + len(tareas), 4 + len(tareas))
    wb.save(destino)
    return destino


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--salida', required=True)
    a = ap.parse_args()
    if os.path.isdir(a.salida):
        shutil.rmtree(a.salida)
    os.makedirs(a.salida)
    print('plan      →', plan(a.salida))
    print('checklist →', checklist(a.salida), '(%d tareas: %s)' % (
        sum(len(t) for _, _, t in DC.FASES), '/'.join(str(len(t)) for _, _, t in DC.FASES)))


if __name__ == '__main__':
    main()
