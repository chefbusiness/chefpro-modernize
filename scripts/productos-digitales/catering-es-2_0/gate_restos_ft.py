#!/usr/bin/env python3
"""
gate_restos_ft.py — Catering ES 2.2 (REVISION-FINAL B1 y m13): en los 11 ficheros de `dl/plan-catering-tematico-eventos/`
  · 0 restos del molde Food Truck («food truck», «food-truck», «plan-negocio-food-truck») en celdas (valores y
    literales de fórmula), docProps de los xlsx y texto, tablas y docProps de los docx;
  · los 8 docx en A4, con autor «AI Chef Pro», sin «python-docx» en los metadatos y con fecha de creación ≥ 2026.
Corre en el VPS. Uso: gate_restos_ft.py [carpeta] · exit 1 si algo falla. Sesión Claude Code, 11-oct-2026.
"""
import os
import re
import sys

import docx
import openpyxl
from docx.shared import Mm

AQUI = os.path.dirname(os.path.abspath(__file__))
DL = os.path.abspath(os.path.join(AQUI, '..', '..', '..', 'astro-site', 'public', 'dl', 'plan-catering-tematico-eventos'))
RX = re.compile(r'food[\s-]?truck|plan-negocio-food-truck', re.I)


def main(carpeta):
    err, n = [], 0
    for f in sorted(os.listdir(carpeta)):
        p = os.path.join(carpeta, f)
        n += 1
        if f.endswith('.xlsx'):
            wb = openpyxl.load_workbook(p)
            for ws in wb.worksheets:
                if RX.search(ws.title):
                    err.append('%s: pestaña «%s»' % (f, ws.title))
                for row in ws.iter_rows():
                    for c in row:
                        if isinstance(c.value, str) and RX.search(c.value):
                            err.append('%s · %s!%s: %s' % (f, ws.title, c.coordinate, c.value[:90]))
            pr = wb.properties
            for campo in ('title', 'subject', 'description', 'keywords', 'category', 'creator', 'lastModifiedBy'):
                v = getattr(pr, campo) or ''
                if RX.search(str(v)):
                    err.append('%s · docProps %s: %s' % (f, campo, v))
        elif f.endswith('.docx'):
            d = docx.Document(p)
            textos = [x.text for x in d.paragraphs]
            for t in d.tables:
                for row in t.rows:
                    textos += [c.text for c in row.cells]
            for s in d.sections:
                for hf in (s.header, s.footer, s.first_page_header, s.first_page_footer):
                    textos += [x.text for x in hf.paragraphs]
            for t in textos:
                if RX.search(t):
                    err.append('%s: «%s»' % (f, t[:90]))
            cp = d.core_properties
            for campo in ('title', 'subject', 'comments', 'keywords', 'category', 'author', 'last_modified_by'):
                v = getattr(cp, campo) or ''
                if RX.search(v) or 'python-docx' in v:
                    err.append('%s · docProps %s: %s' % (f, campo, v))
            if cp.author != 'AI Chef Pro':
                err.append('%s · autor %r' % (f, cp.author))
            if not cp.created or cp.created.year < 2026:
                err.append('%s · fecha de creación %s' % (f, cp.created))
            for s in d.sections:
                if abs(s.page_width - Mm(210)) > Mm(1) or abs(s.page_height - Mm(297)) > Mm(1):
                    err.append('%s · página %.1f × %.1f mm (no A4)' % (f, s.page_width.mm, s.page_height.mm))
                    break
    for e in err:
        print('ERROR', e)
    print('gate restos FT + metadatos: %d ficheros · %s' % (n, 'VERDE' if not err else 'ROJO (%d)' % len(err)))
    return 1 if err else 0


if __name__ == '__main__':
    sys.exit(main(sys.argv[1] if len(sys.argv) > 1 else DL))
