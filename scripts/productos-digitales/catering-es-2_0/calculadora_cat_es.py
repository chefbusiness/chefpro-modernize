#!/usr/bin/env python3
"""
calculadora_cat_es.py — Catering ES 2.0 (SPEC D8b): `calculadora-pricing-catering-tematico-eventos.xlsx`, del coste
al precio por invitado y al presupuesto completo. Copia adaptada de `business-plans-en-4/calculadora_cat.py`:
  · IVA del catering con servicio al 10 % (Ley 37/1992, art. 91.Uno.2.2.º); precio a particulares CON IVA;
  · sin service charge ni sales tax; la propina voluntaria queda fuera del total (nota);
  · anticipo del 40 % y tramos de cancelación de los contratos del kit (particulares 21/7 días + 48 h; empresas
    28/14/7 días), con un selector Particular/Empresa;
  · recargo de desplazamiento por km de ida y vuelta; nº de estaciones temáticas como entrada;
  · margen bruto CALCULADO (nunca tecleado); 3 presupuestos de ejemplo por las mismas fórmulas; autotest pycel.
Hojas `Calculadora` · `Ejemplos` · `Instrucciones` (firma `calc-beta` de `motor.detectar`). A4, protegida sin
contraseña, verdes desbloqueadas, línea de versión y bio de la familia. Corre en el VPS (openpyxl + pycel).

    /root/venv-guias/bin/python calculadora_cat_es.py --salida <dir>      # genera + inject_cache + autotest
    /root/venv-guias/bin/python calculadora_cat_es.py --salida <dir> --autotest

Sesión Claude Code, 11-oct-2026.
"""
import argparse
import math
import os
import re
import sys
from collections import OrderedDict

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill, Protection
from openpyxl.worksheet.datavalidation import DataValidation

AQUI = os.path.dirname(os.path.abspath(__file__))
PD = os.path.dirname(AQUI)
REPO = os.path.abspath(os.path.join(AQUI, '..', '..', '..'))
sys.path.insert(0, os.path.join(PD, 'planes-v2_0'))
sys.path.insert(0, AQUI)
import motor                                                                       # noqa: E402
import tandas_docx as T                                                            # noqa: E402

PID = 'plan-catering-tematico-eventos'
FICHERO = 'calculadora-pricing-catering-tematico-eventos.xlsx'
TITULO = 'Calculadora de precios de catering (del coste al presupuesto)'
H1, H2, H3 = 'Calculadora', 'Ejemplos', 'Instrucciones'
VERDE = PatternFill('solid', fgColor='E8F5E9')
GRIS = PatternFill('solid', fgColor='F2F2F2')
FUENTE = 'Calibri'
ESTILOS = ['Entrega sin servicio', 'Bufé', 'Cóctel', 'Emplatado', 'Estaciones temáticas']
CANALES = ['Particular', 'Empresa']
EUR, ENT, PCT, DEC = '#,##0.00 €', '#,##0', '0%', '#,##0.0'

ENTRADAS = [
    ('canal', 'Cliente', '@', 'Empresa',
     'Particular o Empresa: cambia los tramos de cancelación (contratos del kit) y cómo se presenta el precio.'),
    ('guests', 'Invitados previstos', ENT, 80, 'Número entero. Comida, bebida y desechables se costean con ellos.'),
    ('minimum', 'Mínimo facturable de invitados (contrato)', ENT, 60,
     'Si confirman menos, facturas el mínimo. 0 = sin mínimo.'),
    ('style', 'Estilo de servicio', '@', 'Cóctel',
     'Entrega sin servicio, Bufé, Cóctel, Emplatado o Estaciones temáticas: fija el personal sugerido (tabla de abajo).'),
    ('stations', 'Estaciones temáticas en directo (número)', ENT, 2,
     'Cada estación suma su coste de montaje y un cocinero al personal sugerido.'),
    ('st_cost', 'Coste por estación (gas, montaje y menaje específico)', EUR, 60.0, 'Estimación del kit.'),
    ('food', 'Coste de comida por invitado (escandallo)', EUR, 9.0,
     'Materia prima del menú por invitado. Usa tu escandallo; el valor es una estimación del kit.'),
    ('beverage', 'Coste de bebida por invitado', EUR, 4.0,
     'Vino, cerveza, refrescos y agua que sirves tú. Si el cliente pone la bebida, 0 (y cobra el descorche).'),
    ('disposables', 'Coste de desechables, hielo y consumibles por invitado', EUR, 0.8, 'Estimación del kit.'),
    ('staff', 'Personal del evento (personas)', ENT, 6, 'Camareros, barra, cocineros de estación y coordinación.'),
    ('hours', 'Horas por persona (carga, viaje, montaje, servicio y recogida)', DEC, 7, 'Estimación del kit.'),
    ('wage', 'Coste bruto por hora del personal de eventos', EUR, 11.0,
     'Estimación del kit (1.400 € brutos al mes en 14 pagas ≈ 11 €/h). Copia tu convenio provincial de hostelería.'),
    ('ss', 'Seguridad Social a cargo de la empresa', PCT, 0.33,
     '30,65 % de contingencias, desempleo, FOGASA, formación y MEI, más AT/EP (fijos discontinuos, art. 16 ET).'),
    ('rentals', 'Alquileres repercutidos en el precio (mobiliario, carpas, menaje extra)', EUR, 0.0,
     'Lo que te cobra el alquilador para este evento. Si lo facturas aparte, ese importe va al 21 % de IVA.'),
    ('km', 'Kilómetros de ida (cocina central → espacio)', ENT, 20, 'Se cobran ida y vuelta.'),
    ('per_km', 'Coste por kilómetro de la furgoneta', EUR, 0.40,
     'Estimación del kit (combustible y desgaste): pon tu coste real.'),
    ('margin', 'Margen bruto objetivo sobre el subtotal', PCT, 0.50,
     'Parte del subtotal que queda después de los costes directos de arriba. Estimación del kit.'),
    ('iva', 'IVA del catering con servicio', PCT, 0.10,
     'Ley 37/1992, art. 91.Uno.2.2.º: comida y bebida para consumir en el acto, alcohol incluido.'),
    ('deposit', 'Anticipo para bloquear la fecha (% del total)', PCT, 0.40,
     'El de los contratos del kit: 40 % al firmar. Particulares: el resto 7 días antes; empresas: 30 % siete días '
     'antes y 30 % a 30 días de la factura.'),
]
K = OrderedDict((k, 5 + i) for i, (k, *_r) in enumerate(ENTRADAS))
F_CALC = 5 + len(ENTRADAS) + 1
SALIDAS = [
    ('billable', 'Invitados facturables', ENT, 'max({guests},{minimum})', 'El mayor de previstos y mínimo.'),
    ('c_food', 'Comida, bebida y desechables', EUR, '({food}+{beverage}+{disposables})*{guests}', ''),
    ('c_staff', 'Personal del evento (salario + Seguridad Social)', EUR, '{staff}*{hours}*{wage}*(1+{ss})', ''),
    ('c_st', 'Estaciones temáticas', EUR, '{stations}*{st_cost}', ''),
    ('c_other', 'Alquileres y desplazamiento (ida y vuelta)', EUR, '{rentals}+{km}*2*{per_km}', ''),
    ('c_total', 'Coste directo total', EUR, '{c_food}+{c_staff}+{c_st}+{c_other}', ''),
    ('c_guest', 'Coste directo por invitado facturable', EUR, '{c_total}/{billable}', ''),
    ('price', 'PRECIO POR INVITADO SIN IVA', EUR, 'roundup({c_total}/(1-{margin})/{billable}*2)/2',
     'Redondeado a los 0,50 € de arriba. Es el que va en el presupuesto a empresas («+ IVA»).'),
    ('price_iva', 'PRECIO POR INVITADO CON IVA', EUR, '{price}*(1+{iva})',
     'Es el que ve un particular: en una oferta a consumidores el precio va con impuestos incluidos.'),
    ('subtotal', 'Subtotal sin IVA', EUR, '{price}*{billable}', 'Precio por invitado × invitados facturables.'),
    ('iva_amt', 'IVA', EUR, '{subtotal}*{iva}', ''),
    ('total', 'TOTAL DEL PRESUPUESTO (con IVA)', EUR, '{subtotal}+{iva_amt}', ''),
    ('dep_amt', 'Anticipo para bloquear la fecha', EUR, '{total}*{deposit}', ''),
    ('balance', 'Pendiente después del anticipo', EUR, '{total}-{dep_amt}', ''),
    ('profit', 'Margen bruto en euros', EUR, '{subtotal}-{c_total}', ''),
    ('margin_calc', 'Margen bruto (calculado)', '0.0%', 'if({subtotal}==0,0,{profit}/{subtotal})',
     'Sale de los costes, nunca se teclea. Queda igual o por encima del objetivo por el redondeo del precio.'),
    ('can1', 'Importe retenido si cancela con más de 21 días (empresa: más de 28)', EUR, 'if({canal}=="Particular",0,0)',
     'Particular: se devuelve el anticipo. Empresa: se devuelve el anticipo.'),
    ('can2', 'Importe retenido si cancela entre 21 y 7 días (empresa: entre 28 y 14)', EUR,
     'if({canal}=="Particular",{dep_amt}*0.5,{total}*0.3)',
     'Particular: se retiene el 50 % del anticipo. Empresa: el 30 % del total.'),
    ('can3', 'Importe retenido si cancela con menos de 7 días (empresa: entre 14 y 7)', EUR,
     'if({canal}=="Particular",{dep_amt},{total}*0.6)',
     'Particular: se retiene el anticipo entero. Empresa: el 60 % del total.'),
    ('can4', 'Importe retenido si cancela con menos de 48 horas (empresa: menos de 7 días)', EUR,
     'if({canal}=="Particular",{dep_amt}+{total}*0.5,{total})',
     'Particular: el anticipo más el 50 % del total. Empresa: el 100 % del total.'),
    ('staff_sug', 'Personal sugerido para este estilo (estimación del kit)', ENT, 'sugerido',
     'De la tabla de abajo, más un cocinero por estación.'),
]
S = OrderedDict((k, F_CALC + 1 + i) for i, (k, *_r) in enumerate(SALIDAS))
F_TABLA = F_CALC + 1 + len(SALIDAS) + 1
RATIO = OrderedDict([('Entrega sin servicio', 0), ('Bufé', 20), ('Cóctel', 22), ('Emplatado', 11),
                     ('Estaciones temáticas', 15)])

EJEMPLOS = OrderedDict([
    ('Coffee break en oficina', dict(canal='Empresa', guests=35, minimum=20, style='Entrega sin servicio', stations=0,
                                     st_cost=60.0, food=4.5, beverage=1.5, disposables=0.6, staff=1, hours=3,
                                     wage=11.0, ss=0.33, rentals=0.0, km=15, per_km=0.40, margin=0.50, iva=0.10,
                                     deposit=0.0)),
    ('Cóctel de empresa', dict(canal='Empresa', guests=80, minimum=60, style='Cóctel', stations=2, st_cost=60.0,
                               food=9.0, beverage=4.0, disposables=0.8, staff=6, hours=7, wage=11.0, ss=0.33,
                               rentals=0.0, km=20, per_km=0.40, margin=0.50, iva=0.10, deposit=0.40)),
    ('Boda con estaciones', dict(canal='Particular', guests=110, minimum=90, style='Estaciones temáticas', stations=3,
                                 st_cost=60.0, food=20.0, beverage=7.0, disposables=1.0, staff=10, hours=9,
                                 wage=11.0, ss=0.33, rentals=600.0, km=40, per_km=0.40, margin=0.50, iva=0.10,
                                 deposit=0.40)),
])
EJ_NOTAS = OrderedDict([
    ('Coffee break en oficina', 'Entrega sin servicio: una persona lleva y monta; la empresa paga la factura (sin '
                                'anticipo). Si en la entrega va alcohol o refrescos, zumos y '
                                'gaseosas con azúcares o edulcorantes añadidos, esa parte va al 21 % de IVA.'),
    ('Cóctel de empresa', 'Cóctel de presentación con dos estaciones en directo; presupuesto «+ IVA» y tramos de '
                          'cancelación de empresa.'),
    ('Boda con estaciones', 'Banquete con tres estaciones temáticas y mobiliario alquilado repercutido; precio con IVA '
                            'para el particular y anticipo del 40 %.'),
])


def calcular(e):
    r = OrderedDict()
    r['billable'] = max(e['guests'], e['minimum'])
    r['c_food'] = (e['food'] + e['beverage'] + e['disposables']) * e['guests']
    r['c_staff'] = e['staff'] * e['hours'] * e['wage'] * (1 + e['ss'])
    r['c_st'] = e['stations'] * e['st_cost']
    r['c_other'] = e['rentals'] + e['km'] * 2 * e['per_km']
    r['c_total'] = r['c_food'] + r['c_staff'] + r['c_st'] + r['c_other']
    r['c_guest'] = r['c_total'] / r['billable']
    r['price'] = math.ceil(round(r['c_total'] / (1 - e['margin']) / r['billable'] * 2, 9)) / 2
    r['price_iva'] = r['price'] * (1 + e['iva'])
    r['subtotal'] = r['price'] * r['billable']
    r['iva_amt'] = r['subtotal'] * e['iva']
    r['total'] = r['subtotal'] + r['iva_amt']
    r['dep_amt'] = r['total'] * e['deposit']
    r['balance'] = r['total'] - r['dep_amt']
    r['profit'] = r['subtotal'] - r['c_total']
    r['margin_calc'] = 0 if r['subtotal'] == 0 else r['profit'] / r['subtotal']
    part = e['canal'] == 'Particular'
    r['can1'] = 0.0
    r['can2'] = r['dep_amt'] * 0.5 if part else r['total'] * 0.3
    r['can3'] = r['dep_amt'] if part else r['total'] * 0.6
    r['can4'] = r['dep_amt'] + r['total'] * 0.5 if part else r['total']
    ratio = RATIO[e['style']]
    r['staff_sug'] = (1 if ratio == 0 else math.ceil(e['guests'] / ratio)) + e['stations']
    return r


def a_excel(expr, ref):
    f = expr.replace('max(', 'MAX(').replace('roundup(', 'ROUNDUP(').replace('if(', 'IF(').replace('==', '=')
    f = re.sub(r'\{(\w+)\}', lambda m: ref[m.group(1)], f)
    return f.replace('*2)/2', '*2,0)/2')


def formula_sugerido(ref_estilo, ref_guests, ref_st):
    rng_e = "'%s'!$A$%d:$A$%d" % (H1, F_TABLA + 1, F_TABLA + len(RATIO))
    rng_r = "'%s'!$B$%d:$B$%d" % (H1, F_TABLA + 1, F_TABLA + len(RATIO))
    rat = 'INDEX(%s,MATCH(%s,%s,0))' % (rng_r, ref_estilo, rng_e)
    return '=IFERROR(IF(%s=0,1,ROUNDUP(%s/%s,0))+%s,"—")' % (rat, ref_guests, rat, ref_st)


def _f(ws, coord, v, b=False, sz=11, i=False, fmt=None, wrap=False, fill=None, locked=True):
    c = ws[coord]
    c.value = v
    c.font = Font(name=FUENTE, size=sz, bold=b, italic=i)
    if fmt:
        c.number_format = fmt
    c.alignment = Alignment(wrap_text=wrap, vertical='top')
    if fill:
        c.fill = fill
    c.protection = Protection(locked=locked)
    return c


def _imprimir(ws, horizontal=False):
    ws.page_setup.paperSize = 9                                                   # A4
    ws.page_setup.orientation = 'landscape' if horizontal else 'portrait'
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.oddFooter.center.text = 'aichef.pro/%s · página &P de &N' % PID
    ws.print_options.horizontalCentered = True
    ws.protection.sheet = True                                                    # sin contraseña
    ws.sheet_view.showGridLines = False


def _dv(ws, tipo, celdas, f1=None, f2=None, op=None, titulo='Valor no válido', msg=''):
    dv = DataValidation(type=tipo, formula1=f1, formula2=f2, operator=op, allow_blank=False, showErrorMessage=True,
                        errorTitle=titulo, error=msg)
    ws.add_data_validation(dv)
    for c in celdas:
        dv.add(c)


def hoja_calculadora(wb):
    ws = wb.active
    ws.title = H1
    for col, w in (('A', 62), ('B', 18), ('C', 80)):
        ws.column_dimensions[col].width = w
    _f(ws, 'A1', 'CALCULADORA DE PRECIOS DE CATERING — DEL COSTE AL PRESUPUESTO', b=True, sz=14)
    _f(ws, 'A2', 'Escribe sólo en las celdas VERDES: todo lo demás se calcula. Importes sin IVA salvo donde la fila '
                 'diga lo contrario. Los valores por defecto son estimaciones del kit: cámbialos por los tuyos.',
       wrap=True)
    ws.merge_cells('A2:C2')
    ws.row_dimensions[2].height = 32
    for col, t in zip('ABC', ('DATOS DEL EVENTO', 'Valor', 'Notas')):
        _f(ws, '%s4' % col, t, b=True, fill=GRIS)
    for k, rot, fmt, val, nota in ENTRADAS:
        f = K[k]
        _f(ws, 'A%d' % f, rot)
        _f(ws, 'B%d' % f, val, fmt=fmt if fmt != '@' else None, fill=VERDE, locked=False)
        _f(ws, 'C%d' % f, nota, wrap=True)
    for col, t in zip('ABC', ('CÁLCULO', 'Valor', 'Notas')):
        _f(ws, '%s%d' % (col, F_CALC), t, b=True, fill=GRIS)
    ref = dict((k, '$B$%d' % f) for k, f in K.items())
    ref.update((k, '$B$%d' % f) for k, f in S.items())
    for k, rot, fmt, expr, nota in SALIDAS:
        f = S[k]
        fx = (formula_sugerido(ref['style'], ref['guests'], ref['stations']) if expr == 'sugerido'
              else '=' + a_excel(expr, ref))
        neg = k in ('price', 'price_iva', 'total', 'margin_calc')
        _f(ws, 'A%d' % f, rot, b=neg)
        _f(ws, 'B%d' % f, fx, fmt=fmt, b=neg)
        if nota:
            _f(ws, 'C%d' % f, nota, wrap=True)
    for col, t in zip('ABC', ('INVITADOS POR CAMARERO SEGÚN EL ESTILO (estimación del kit, editable)', 'Invitados',
                              'Notas')):
        _f(ws, '%s%d' % (col, F_TABLA), t, b=True, fill=GRIS)
    notas = {'Entrega sin servicio': '0 = sin camareros: una persona lleva y monta.',
             'Emplatado': 'Un camarero por cada 10-12 invitados sentados.',
             'Cóctel': 'Un camarero por cada 20-25 invitados.'}
    for i, (est, n) in enumerate(RATIO.items(), 1):
        _f(ws, 'A%d' % (F_TABLA + i), est)
        _f(ws, 'B%d' % (F_TABLA + i), n, fmt=ENT, fill=VERDE, locked=False)
        if est in notas:
            _f(ws, 'C%d' % (F_TABLA + i), notas[est], wrap=True)
    _dv(ws, 'list', [ref['style'].replace('$', '')], f1='"%s"' % ','.join(ESTILOS), titulo='Estilo no válido',
        msg='Elige un estilo de servicio de la lista.')
    _dv(ws, 'list', [ref['canal'].replace('$', '')], f1='"%s"' % ','.join(CANALES), titulo='Cliente no válido',
        msg='Elige Particular o Empresa.')
    _dv(ws, 'whole', ['B%d' % K['guests']], f1='1', f2='100000', op='between', titulo='Número no válido',
        msg='Escribe un número entero de 1 o más.')
    _dv(ws, 'whole', ['B%d' % K[k] for k in ('minimum', 'staff', 'stations', 'km')] +
        ['B%d' % (F_TABLA + i) for i in range(1, len(RATIO) + 1)], f1='0', f2='100000', op='between',
        titulo='Número no válido', msg='Escribe un número entero de 0 o más.')
    _dv(ws, 'decimal', ['B%d' % K[k] for k in ('st_cost', 'food', 'beverage', 'disposables', 'hours', 'wage',
                                              'rentals', 'per_km')], f1='0', op='greaterThanOrEqual',
        titulo='Importe no válido', msg='Escribe un número mayor o igual que 0.')
    _dv(ws, 'decimal', ['B%d' % K[k] for k in ('ss', 'iva', 'deposit')], f1='0', f2='1', op='between',
        titulo='Porcentaje no válido', msg='Escribe un porcentaje entre 0 y 1 (0,10 = 10 %).')
    _dv(ws, 'decimal', ['B%d' % K['margin']], f1='0', f2='0.95', op='between', titulo='Porcentaje no válido',
        msg='Escribe un margen objetivo entre 0 y 95 % (0,50 = 50 %).')
    ws.freeze_panes = 'A5'
    _imprimir(ws)


def hoja_ejemplos(wb):
    ws = wb.create_sheet(H2)
    ws.column_dimensions['A'].width = 62
    for col in 'BCD':
        ws.column_dimensions[col].width = 24
    _f(ws, 'A1', '3 PRESUPUESTOS DE EJEMPLO — con las mismas fórmulas de la hoja «%s»' % H1, b=True, sz=14)
    _f(ws, 'A2', 'Estimaciones del kit para ilustrar el cálculo, no precios de mercado: cambia los datos en «%s» '
                 'para presupuestar tus eventos.' % H1, wrap=True)
    ws.merge_cells('A2:D2')
    ws.row_dimensions[2].height = 32
    _f(ws, 'A4', 'DATOS DEL EVENTO', b=True, fill=GRIS)
    for j, nombre in enumerate(EJEMPLOS):
        _f(ws, '%s4' % 'BCD'[j], nombre, b=True, fill=GRIS, wrap=True)
    for k, rot, fmt, val, nota in ENTRADAS:
        _f(ws, 'A%d' % K[k], rot)
        for j, e in enumerate(EJEMPLOS.values()):
            _f(ws, '%s%d' % ('BCD'[j], K[k]), e[k], fmt=fmt if fmt != '@' else None)
    _f(ws, 'A%d' % F_CALC, 'RESULTADO', b=True, fill=GRIS)
    for j in range(3):
        _f(ws, '%s%d' % ('BCD'[j], F_CALC), list(EJEMPLOS)[j], b=True, fill=GRIS, wrap=True)
    for k, rot, fmt, expr, nota in SALIDAS:
        neg = k in ('price', 'price_iva', 'total', 'margin_calc')
        _f(ws, 'A%d' % S[k], rot, b=neg)
        for j in range(3):
            col = 'BCD'[j]
            ref = dict((kk, '$%s$%d' % (col, f)) for kk, f in K.items())
            ref.update((kk, '$%s$%d' % (col, f)) for kk, f in S.items())
            fx = (formula_sugerido(ref['style'], ref['guests'], ref['stations']) if expr == 'sugerido'
                  else '=' + a_excel(expr, ref))
            _f(ws, '%s%d' % (col, S[k]), fx, fmt=fmt, b=neg)
    f = S[SALIDAS[-1][0]] + 2
    _f(ws, 'A%d' % f, 'QUÉ ENSEÑA CADA EJEMPLO', b=True, fill=GRIS)
    for j, (nombre, nota) in enumerate(EJ_NOTAS.items()):
        _f(ws, 'A%d' % (f + 1 + j), nombre, b=True)
        _f(ws, 'B%d' % (f + 1 + j), nota, wrap=True)
        ws.merge_cells('B%d:D%d' % (f + 1 + j, f + 1 + j))
        ws.row_dimensions[f + 1 + j].height = 32
    ws.freeze_panes = 'B5'
    _imprimir(ws, horizontal=True)


INSTRUCCIONES = [
    ('INSTRUCCIONES DE USO — Calculadora de precios de catering', True),
    ('Convierte tus costes en un precio por invitado y en un presupuesto completo: subtotal, IVA, total, anticipo y '
     'lo que retienes si el cliente cancela. Sólo se escribe en las celdas VERDES; el margen se calcula.', False),
    ('1. Elige el cliente (Particular o Empresa), los invitados previstos y el mínimo facturable del contrato.', False),
    ('2. Elige el estilo de servicio y el número de estaciones temáticas: el personal sugerido sale de la tabla del '
     'final de la primera hoja (invitados por camarero, editable) más un cocinero por estación.', False),
    ('3. Escribe el coste de comida por invitado de tu escandallo, la bebida que sirves y los desechables.', False),
    ('4. Escribe el personal del evento (personas, horas por persona, coste bruto por hora y Seguridad Social), los '
     'alquileres que repercutes y los kilómetros de ida: el desplazamiento se cobra ida y vuelta.', False),
    ('5. Fija el margen bruto objetivo. El precio por invitado es el coste directo dividido entre (1 - margen), '
     'redondeado a los 0,50 € de arriba: el margen calculado queda igual o por encima del objetivo.', False),
    ('6. A una empresa se le presupuesta el precio SIN IVA («+ IVA»); a un particular, el precio CON IVA, porque en '
     'una oferta a consumidores el precio va con los impuestos incluidos.', False),
    ('IVA, PROPINAS Y SERVICIOS APARTE', True),
    ('El catering con servicio va al 10 % de IVA, alcohol incluido (Ley 37/1992, art. 91.Uno.2.2.º). En una '
     'entrega sin servicio, la comida sigue al 10 % pero el alcohol y los refrescos, zumos y gaseosas con azúcares o edulcorantes añadidos van al 21 %. '
     'El menaje, la mantelería y la cristalería que facturas junto con el catering siguen al servicio principal; si '
     'facturas aparte mobiliario, carpas o decoración, esa línea va al 21 %.', False),
    ('La propina voluntaria es decisión del cliente y es del equipo: no entra en el precio ni en el total.', False),
    ('ANTICIPO Y CANCELACIÓN', True),
    ('Los importes de cancelación siguen los contratos del kit. Particulares: más de 21 días, se devuelve el '
     'anticipo; entre 21 y 7 días, se retiene la mitad; menos de 7 días, el anticipo entero; menos de 48 horas, el '
     'anticipo más el 50 % del total. Empresas: más de 28 días, se devuelve el anticipo; 28-14 días, el 30 % del '
     'total; 14-7 días, el 60 %; menos de 7 días, el 100 %. En un servicio de comida con fecha no hay derecho de '
     'desistimiento (art. 103.l del TRLGDCU): por eso la política tiene que estar firmada.', False),
    ('ESTIMACIONES DEL KIT', True),
    ('Los valores por defecto y los tres ejemplos son estimaciones para ilustrar el cálculo, no precios de mercado. '
     'Cámbialos por tu escandallo, tus salarios, tus alquileres y tus kilómetros antes de presupuestar.', False),
]


def hoja_instrucciones(wb):
    ws = wb.create_sheet(H3)
    ws.column_dimensions['A'].width = 110
    f = 1
    for t, b in INSTRUCCIONES:
        _f(ws, 'A%d' % f, t, b=b, sz=13 if f == 1 else 11, wrap=True)
        if not b:
            ws.row_dimensions[f].height = max(16, 15 * math.ceil(len(t) / 120.0))
        f += 2 if b and f > 1 else 1
    f += 1
    for t in ('Para editar una celda que no esté en verde: Revisar → Desproteger hoja (no tiene contraseña).',
              'Más plantillas y kits del catálogo en aichef.pro/productos-digitales', motor.BIO_LINE,
              motor.VERSION_FMT.format(pid=PID)):
        _f(ws, 'A%d' % f, t, wrap=True)
        f += 1
    _imprimir(ws)


def generar(destino):
    wb = openpyxl.Workbook()
    hoja_calculadora(wb)
    hoja_ejemplos(wb)
    hoja_instrucciones(wb)
    pr = wb.properties
    pr.title = '%s · v%s' % (TITULO, motor.VERSION)
    pr.subject = 'Plan de Negocio: Catering Temático de Eventos'
    pr.creator = pr.lastModifiedBy = motor.CREATOR
    os.makedirs(os.path.dirname(destino), exist_ok=True)
    wb.save(destino)
    sys.path.insert(0, PD)
    import inject_cache as IC                                                      # noqa: E402
    nf, inj, fail = IC.inject(destino)
    print('%s: formulas=%d cache_inyectado=%d fallos_pycel=%d' % (os.path.basename(destino), nf, inj, fail))
    return fail


def autotest(destino):
    from pycel import ExcelCompiler
    err, n = [], 0
    xl = ExcelCompiler(filename=destino)
    wv = openpyxl.load_workbook(destino, data_only=True)
    wf = openpyxl.load_workbook(destino)
    casos = [(H2, 'BCD'[j], e, nombre) for j, (nombre, e) in enumerate(EJEMPLOS.items())]
    casos.append((H1, 'B', OrderedDict((k, v) for k, _r, _f2, v, _n in ENTRADAS), 'hoja principal'))
    for hoja, col, e, nombre in casos:
        esperado = calcular(e)
        for k, v in esperado.items():
            c = '%s%d' % (col, S[k])
            n += 1
            for lab, x in (('pycel', xl.evaluate("'%s'!%s" % (hoja, c))), ('caché', wv[hoja][c].value)):
                if not isinstance(x, (int, float)) or abs(float(x) - v) > 0.005:
                    err.append('%s · %s %s (%s): %s %r ≠ esperado %r' % (nombre, hoja, c, k, lab, x, v))
        if esperado['margin_calc'] + 1e-9 < e['margin']:
            err.append('%s: margen calculado %.4f < objetivo %.4f' % (nombre, esperado['margin_calc'], e['margin']))
    e = dict(EJEMPLOS['Cóctel de empresa'])
    if calcular(dict(e, canal='Particular'))['can4'] == calcular(dict(e, canal='Empresa'))['can4']:
        err.append('el selector Particular/Empresa no cambia la cancelación')
    if motor.detectar(wf, FICHERO)['molde'] != 'calc-beta':
        err.append('motor.detectar no la reconoce como calc-beta')
    textos = []
    for ws in wf.worksheets:
        if ws.page_setup.paperSize != 9 or not ws.protection.sheet:
            err.append('%s: A4 y protección' % ws.title)
        for row in ws.iter_rows():
            for c in row:
                v = c.value
                if isinstance(v, str) and c.data_type != 'f':
                    textos.append((ws.title, c.coordinate, v))
                if c.data_type == 'f' and wv[ws.title][c.coordinate].value is None:
                    err.append('%s!%s: fórmula sin caché' % (ws.title, c.coordinate))
                if c.fill.fgColor.rgb in ('00E8F5E9', 'FFE8F5E9') and c.protection.locked:
                    err.append('%s!%s: verde bloqueada' % (ws.title, c.coordinate))
    todo = '\n'.join(t for _h, _c, t in textos)
    for rx in T.PROHIBIDOS:
        if re.search(rx, todo, re.I):
            err.append('prohibido /%s/' % rx)
    if re.search(r'service charge|sales tax|\$', todo, re.I):
        err.append('restos del EN (service charge / sales tax / $)')
    if motor.VERSION_FMT.format(pid=PID) not in todo or motor.BIO_LINE not in todo:
        err.append('falta la línea de versión o la bio')
    if wf.sheetnames != [H1, H2, H3]:
        err.append('pestañas %r' % wf.sheetnames)
    for x in err:
        print('ERROR', x)
    print('autotest calculadora ES: %d comprobaciones (pycel + caché) sobre 3 ejemplos + la hoja principal · %s' % (
        n * 2, 'VERDE' if not err else 'ROJO (%d errores)' % len(err)))
    for nombre, e in EJEMPLOS.items():
        r = calcular(e)
        print('  %-24s %3d inv · %6.2f €/inv sin IVA · %6.2f con IVA · total %9.2f · anticipo %8.2f · margen %.1f %% · '
              'personal sugerido %d' % (nombre, r['billable'], r['price'], r['price_iva'], r['total'], r['dep_amt'],
                                        r['margin_calc'] * 100, r['staff_sug']))
    return 1 if err else 0


if __name__ == '__main__':
    ap = argparse.ArgumentParser()
    ap.add_argument('--salida', required=True, help='carpeta de salida')
    ap.add_argument('--autotest', action='store_true')
    a = ap.parse_args()
    destino = os.path.join(a.salida, FICHERO)
    if not a.autotest and generar(destino):
        sys.exit('ERROR: fallos de pycel en inject_cache')
    sys.exit(autotest(destino))
