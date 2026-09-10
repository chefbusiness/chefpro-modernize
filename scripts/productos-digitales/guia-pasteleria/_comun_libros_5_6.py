#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
_comun_libros_5_6.py — utillaje compartido por los generadores de los libros
**5** (`plan-financiero-3-anos-pasteleria`) y **6** (`checklist-legal-y-licencias`)
de «Cómo Montar una Pastelería» (SPEC §2.2, filas 5 y 6).

NO es un motor paralelo: `guias-v2_0/motor.py` sigue poniendo las verdes, las
fórmulas registradas, los semáforos, las validaciones y la protección. Aquí sólo
vive lo que estos dos libros repiten y el motor no puede dar:

  1. **`VERSION_LINE` propio.** `motor.version_line()` devuelve «Versión 2.0 ·
     agosto 2026 · …» (motor.py:203-215), que es la línea de la familia
     `guias-v2_0`. Este producto NACE en 1.0 y en septiembre de 2026. Ojo con el
     libro 5: es HÍBRIDO (molde `planes-v2_0` motor 2.2 + hoja propia) y manda
     esta línea, no el literal de versión del motor de planes (SPEC §9).
  2. **`dv_rango()`.** `motor.dv_lista()` construye la validación con una cadena
     de opciones separadas por comas y la SPEC §2.2 exige que TODA lista vaya
     contra un RANGO de la propia hoja.
  3. **`nota_celda()`**: la nota legal de `datos_ejemplo.nota_legal(id)` como
     comentario de celda, contada para el resumen.
  4. **`cerrar()`**: guardar, `inject_cache.py`, verificación `data_only` de
     TODAS las fórmulas registradas, censo de verdes vacías (D21), barrido de
     funciones prohibidas y `build/mapa-<libro>.json` en el formato que consume
     el guion (`etiqueta -> {ref, valor, tipo}`).

Los otros dos pares de constructores tienen sus propios `_comun_*`: este fichero
no los toca ni depende de ellos.

Via: Claude Code
"""
import json
import os
import subprocess
import sys

import openpyxl
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.page import PageMargins
from openpyxl.worksheet.properties import PageSetupProperties

AQUI = os.path.dirname(os.path.abspath(__file__))
RAIZ = os.path.normpath(os.path.join(AQUI, '..'))
if os.path.join(RAIZ, 'guias-v2_0') not in sys.path:
    sys.path.insert(0, os.path.join(RAIZ, 'guias-v2_0'))
if AQUI not in sys.path:
    sys.path.insert(0, AQUI)

import motor                                                   # noqa: E402
import datos_ejemplo as D                                      # noqa: E402
import _comun_pasteleria as CP                                 # noqa: E402

motor.CTX['producto'] = 'guia-pasteleria-obrador'

PID = 'guia-pasteleria-obrador'
PRODUCTO = 'Cómo Montar una Pastelería'
SUBTITULO = 'AI Chef Pro · aichef.pro — ' + PRODUCTO
SUBJECT = PRODUCTO + ' · Versión 1.0 · septiembre 2026'

VERSION_LINE = ('Versión 1.0 · septiembre 2026 · aichef.pro/' + PID
                + ' · info@aichef.pro')
#: `motor.BIO_LINE` lleva el sufijo « · johnguerrero.es» de la familia
#: `guias-v2_0`; el literal de la SPEC de este producto termina en punto, sin
#: sufijo (auditoría 2026-09-10, hallazgo A4) — mismo texto que
#: `_comun_pasteleria.BIO`, para que los ocho libros publiquen la misma frase.
BIO = ('Diseñado por John Guerrero — chef y consultor gastronómico desde 2010, '
       'en cocina desde los 17 años.')
assert motor.RX_BIO.search(BIO)
#: `motor.NOTA_DESPROTEGER` lleva una FLECHA (U+2192) fuera de cp1252
#: (hallazgo A5); misma frase, con «>» en vez de flecha.
NOTA_DESPROTEGER = ('Para editar la estructura o una celda que no esté en '
                    'verde: Revisar > Desproteger hoja (no tiene contraseña).')
assert motor.RX_DESPROTEGER.match(NOTA_DESPROTEGER)
NOTA_VERDES = motor.NOTA_VERDES
PIE = 'AI Chef Pro · aichef.pro · Página &P de &N'

EUR = motor.FMT_EUR
PCT = motor.FMT_PCT
PCT2 = '0.00%'
ENT = motor.FMT_ENT
FECHA = motor.FMT_FECHA
DEC = '#,##0.00'
DEC1 = '#,##0.0'

ORO = 'FFD700'
GRIS = 'F2F2F2'
GRIS_TXT = '808080'
CREMA = 'FFF6DC'
CABECERA = '2D2D2D'
AZUL = 'DCE6F1'

ETIQUETA_LISTAS = ('Listas de validación (no las borres: de aquí salen los '
                   'desplegables)')

#: (hoja, coord, id) de cada nota legal puesta por el generador en curso.
NOTAS_LEGALES = []
#: (hoja, coord, etiqueta, valor por defecto) de cada celda verde (D21).
VERDES = []

_FINO = Side(style='thin', color='BFBFBF')
BORDE = Border(left=_FINO, right=_FINO, top=_FINO, bottom=_FINO)


# --------------------------------------------------------------------------
# Formato
# --------------------------------------------------------------------------
def anchos(ws, mapa):
    for letra, ancho in mapa.items():
        ws.column_dimensions[letra].width = ancho


def encabezar(ws, titulo, nota_txt=None, col_fin='F'):
    motor.val(ws, 'A1', CP.titulo_hoja(titulo))
    ws['A1'].font = Font(bold=True, size=16, color=ORO)
    ws.row_dimensions[1].height = 26
    motor.val(ws, 'A2', SUBTITULO)
    ws['A2'].font = Font(size=9, color=GRIS_TXT)
    if nota_txt:
        ws.merge_cells('A3:%s3' % col_fin)
        motor.val(ws, 'A3', nota_txt, wrap=True)
        ws['A3'].font = Font(italic=True, size=9)
        ws.row_dimensions[3].height = 28


def cabecera(ws, fila, columnas, altura=32):
    """`columnas` = [(letra, texto), ...]."""
    for letra, texto in columnas:
        c = ws[letra + str(fila)]
        c.value = texto
        c.fill = PatternFill('solid', fgColor=CABECERA)
        c.font = Font(bold=True, color='FFFFFF', size=9)
        c.alignment = Alignment(horizontal='center', vertical='center',
                                wrap_text=True)
    ws.row_dimensions[fila].height = altura


def seccion(ws, coord, texto):
    motor.val(ws, coord, texto)
    ws[coord].font = Font(bold=True, size=11, color=ORO)
    ws[coord].fill = PatternFill('solid', fgColor=CABECERA)
    return ws[coord]


def total_fila(ws, fila, letras):
    for letra in letras:
        c = ws[letra + str(fila)]
        c.font = Font(bold=True, size=c.font.size)
        c.fill = PatternFill('solid', fgColor=CREMA)


def nota(ws, coord, texto, tam=9):
    motor.val(ws, coord, texto, wrap=True)
    ws[coord].font = Font(italic=True, size=tam)
    return ws[coord]


def parrafo(ws, fila, texto, col_ini='A', col_fin='F', alto=30, tam=9,
            italica=True):
    ws.merge_cells('%s%d:%s%d' % (col_ini, fila, col_fin, fila))
    motor.val(ws, '%s%d' % (col_ini, fila), texto, wrap=True)
    ws['%s%d' % (col_ini, fila)].font = Font(italic=italica, size=tam)
    ws.row_dimensions[fila].height = alto


def pagina(ws, apaisado=True, titulos=None, area=None):
    ws.page_setup.paperSize = 9                        # A4
    ws.page_setup.orientation = 'landscape' if apaisado else 'portrait'
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    if ws.sheet_properties.pageSetUpPr is None:
        ws.sheet_properties.pageSetUpPr = PageSetupProperties()
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_margins = PageMargins(left=0.39, right=0.39, top=0.55,
                                  bottom=0.55, header=0.3, footer=0.3)
    ws.oddFooter.center.text = PIE
    ws.oddFooter.center.size = 8
    if titulos:
        ws.print_title_rows = titulos
    if area:
        ws.print_area = area


# --------------------------------------------------------------------------
# Entradas, notas y desplegables
# --------------------------------------------------------------------------
def entrada(ws, coord, valor, fmt=None, etiqueta='', bold=None, align=None):
    """Celda VERDE de entrada. Registra su valor por defecto (D21: ninguna
    celda verde se queda vacía por depender de otro producto)."""
    cel = motor.val(ws, coord, valor, fmt=fmt, verde_=True, bold=bold,
                    align=align)
    VERDES.append((ws.title, coord, etiqueta, valor))
    return cel


def nota_celda(ws, coord, id_pa, extra=None):
    """Nota legal de `id_pa` como comentario de celda (regla de la SPEC §2.2:
    «Verificado el 10-09-2026 · norma y artículo · URL»)."""
    texto = D.nota_legal(id_pa)
    if not texto:
        raise SystemExit('nota_legal(%r) no devuelve nada: gate_legal() '
                         'tendría que haber abortado antes.' % id_pa)
    if extra:
        texto = texto + '\n' + extra
    ws[coord].comment = Comment(texto, 'AI Chef Pro', height=130, width=460)
    NOTAS_LEGALES.append((ws.title, coord, id_pa))
    return texto


def comentario(ws, coord, texto):
    """Comentario NO legal (criterio propio, supuesto declarado)."""
    ws[coord].comment = Comment(texto, 'AI Chef Pro', height=110, width=440)


def dv_rango(ws, coords, ref, titulo, mensaje, prompt=None, blank=False):
    """Desplegable cuyo origen es un RANGO de la propia hoja (SPEC §2.2:
    DV contra rango, nunca lista con comas, y `showErrorMessage=True`)."""
    if not coords:
        return None
    dv = DataValidation(type='list', formula1=ref, allow_blank=blank,
                        showErrorMessage=True, errorTitle=titulo,
                        error=mensaje)
    if prompt:
        dv.showInputMessage = True
        dv.promptTitle = titulo
        dv.prompt = prompt
    ws.add_data_validation(dv)
    for c in coords:
        dv.add(c)
    return dv


def bloque_listas(ws, fila, listas, col='A'):
    """Escribe al pie de la hoja los vocabularios de los desplegables.

    `listas` = [(nombre, [opción, ...]), ...]. Devuelve {nombre: '$X$a:$X$b'}.
    """
    motor.val(ws, col + str(fila), ETIQUETA_LISTAS)
    ws[col + str(fila)].font = Font(italic=True, bold=True, size=9,
                                    color=GRIS_TXT)
    fila += 1
    refs = {}
    letra = col
    for nombre, opciones in listas:
        motor.val(ws, letra + str(fila), nombre)
        ws[letra + str(fila)].font = Font(bold=True, size=9)
        ini = fila + 1
        for i, op in enumerate(opciones):
            motor.val(ws, letra + str(ini + i), op)
            ws[letra + str(ini + i)].font = Font(size=9)
        refs[nombre] = "'%s'!$%s$%d:$%s$%d" % (ws.title, letra, ini, letra,
                                               ini + len(opciones) - 1)
        fila = ini + len(opciones) + 1
    return refs, fila


def pie(ws, fila, col='A', col_fin='C'):
    """Bio anclada + línea de versión, al pie de la hoja de Instrucciones."""
    ws.merge_cells('%s%d:%s%d' % (col, fila, col_fin, fila))
    motor.val(ws, col + str(fila), BIO, wrap=True)
    ws[col + str(fila)].font = Font(size=9)
    ws.merge_cells('%s%d:%s%d' % (col, fila + 1, col_fin, fila + 1))
    motor.val(ws, col + str(fila + 1), VERSION_LINE)
    ws[col + str(fila + 1)].font = Font(size=9, color=GRIS_TXT)
    return fila + 2


# --------------------------------------------------------------------------
# Cierre: guardar, cachear, verificar y publicar el mapa
# --------------------------------------------------------------------------
PROHIBIDAS = ('INDIRECT(', 'COUNTA(', 'PMT(', 'OFFSET(', 'XLOOKUP(', 'LET(',
              'LAMBDA(', 'RANK(', 'NETWORKDAYS(', 'IRR(', 'MINIFS(',
              'MAXIFS(', 'TEXTJOIN(')


def cerrar(wb, nombre, titulo, mapa_celdas, notas_mapa=None,
           blancos_ok=None):
    """Guarda, inyecta caché, verifica `data_only` y escribe el mapa.

    `mapa_celdas` = [(etiqueta, hoja, coord, tipo), ...] con `tipo` en
    {'entrada', 'salida', 'parametro'}. Devuelve el resumen del libro.

    `blancos_ok` = {(hoja, coord): motivo} para las MUY pocas fórmulas cuyo
    resultado correcto es la cadena vacía (la regla de la familia: «sin dato» =
    `""`, nunca `0`). `inject_cache` no escribe caché para una cadena vacía y
    openpyxl devuelve `None` tanto para «sin caché» como para «caché vacía», así
    que no se pueden distinguir aquí: se declaran una a una, con su motivo, y el
    `demo()` del generador comprueba con pycel que valen `''` y no un error.
    """
    blancos_ok = blancos_ok or {}
    wb.properties.creator = 'AI Chef Pro'
    wb.properties.lastModifiedBy = 'AI Chef Pro'
    wb.properties.title = titulo
    wb.properties.subject = SUBJECT

    for ws in wb.worksheets:
        motor.retirar_verde_de_calculadas(ws)
        motor.proteger(ws)

    destino = os.path.join(AQUI, 'build')
    if not os.path.isdir(destino):
        os.makedirs(destino)
    ruta = os.path.join(destino, nombre + '.xlsx')
    wb.save(ruta)

    # --- funciones prohibidas (antes de gastar un segundo de pycel) --------
    usos = []
    for hoja, coord, formula in motor.REGISTRO:
        arriba = formula.upper().replace(' ', '')
        for p in PROHIBIDAS:
            if p in arriba:
                usos.append('%s!%s usa %s' % (hoja, coord, p[:-1]))
    if usos:
        raise SystemExit('Funciones PROHIBIDAS:\n  ' + '\n  '.join(usos))

    # --- inject_cache ------------------------------------------------------
    res = subprocess.run(
        [sys.executable, os.path.join(RAIZ, 'inject_cache.py'), ruta],
        stdout=subprocess.PIPE, stderr=subprocess.STDOUT)
    salida_cache = res.stdout.decode('utf-8', 'replace')
    if res.returncode != 0:
        raise SystemExit('inject_cache falló:\n' + salida_cache)

    # --- verificación data_only de TODAS las fórmulas registradas ---------
    wbf = openpyxl.load_workbook(ruta)
    wbv = openpyxl.load_workbook(ruta, data_only=True)
    sin_valor, con_error = [], []
    for hoja, coord, formula in motor.REGISTRO:
        v = wbv[hoja][coord].value
        if v is None:
            if (hoja, coord) in blancos_ok:
                continue
            sin_valor.append('%s!%s  %s' % (hoja, coord, formula[:80]))
        elif isinstance(v, str) and v.startswith('#'):
            con_error.append('%s!%s  %s -> %s' % (hoja, coord, formula[:60], v))
    if sin_valor or con_error:
        raise SystemExit(
            'VERIFICACIÓN data_only FALLIDA en %s\n  sin valor (%d):\n   %s\n'
            '  con error (%d):\n   %s'
            % (nombre, len(sin_valor), '\n   '.join(sin_valor[:40]),
               len(con_error), '\n   '.join(con_error[:40])))

    # --- gate de totales: cada `=SUM(rango)` tiene que cachear lo que suma --
    # (hallazgos A1/A3, 10-09-2026: `_comun_libros_3_4.cerrar()` ya lo tenía;
    # aquí no se corría y publicó dos ceros donde había 100 h y 40 h)
    incoherentes = CP.gate_sum_rango(wbv)
    if incoherentes:
        raise SystemExit('TOTALES con caché incoherente en %s:\n  %s'
                         % (nombre, '\n  '.join(incoherentes)))

    # --- barrido WinAnsi (cp1252), hallazgo A5 -----------------------------
    CP.barrer_cp1252(wbf)

    # --- celdas verdes vacías (D21) ---------------------------------------
    verdes_vacias, n_verdes = [], 0
    for ws in wbf.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if c.__class__.__name__ == 'MergedCell':
                    continue
                if motor.es_verde(c):
                    n_verdes += 1
                    if c.value is None or (isinstance(c.value, str)
                                           and not c.value.strip()):
                        verdes_vacias.append('%s!%s' % (ws.title, c.coordinate))

    # --- mapa de celdas citables ------------------------------------------
    mapa = {}
    for etiqueta, hoja, coord, tipo in mapa_celdas:
        if hoja not in wbv.sheetnames:
            raise SystemExit('El mapa cita la hoja inexistente %r' % hoja)
        v = wbv[hoja][coord].value
        if v is None or (isinstance(v, str) and not v.strip()):
            raise SystemExit('El mapa cita %s!%s («%s») y está VACÍA'
                             % (hoja, coord, etiqueta))
        if etiqueta in mapa:
            raise SystemExit('Etiqueta de mapa repetida: %r' % etiqueta)
        if hasattr(v, 'isoformat'):        # fecha o fecha-hora cacheada
            v = v.isoformat()[:10]
        mapa[etiqueta] = {'ref': '%s.xlsx!%s!%s' % (nombre, hoja, coord),
                          'valor': v, 'tipo': tipo}
    if notas_mapa:
        # Hallazgo A9 (2026-09-10): mismo fix que `_comun_libros_3_4.cerrar()`:
        # `ref` con forma `libro.xlsx!Hoja!Celda`, para que `ref.split('!')`
        # no reviente en la mitad de los mapas del producto.
        mapa['_meta'] = {'ref': '%s.xlsx!%s!A1' % (nombre, wb.worksheets[0].title),
                         'valor': notas_mapa, 'tipo': 'nota'}
    with open(os.path.join(destino, 'mapa-' + nombre + '.json'), 'w') as fh:
        fh.write(json.dumps(mapa, ensure_ascii=False, indent=1))

    return {
        'ruta': ruta,
        'hojas': len(wbf.worksheets),
        'formulas': len(motor.REGISTRO),
        'verdes': n_verdes,
        'verdes_vacias': verdes_vacias,
        'notas_legales': len(NOTAS_LEGALES),
        'mapa': len([k for k in mapa if not k.startswith('_')]),
        'blancos_ok': len(blancos_ok),
        'cache': (salida_cache.strip().splitlines() or [''])[-1],
    }


def compilador(ruta):
    """`pycel.ExcelCompiler` sobre el fichero ya escrito, para los `demo()`."""
    import logging
    logging.disable(logging.CRITICAL)
    from pycel import ExcelCompiler
    return ExcelCompiler(ruta)


def resumen(nombre, res, demos):
    ok = all(d[1] for d in demos)
    print('')
    print('=== %s ===' % nombre)
    print('  fichero        : %s' % res['ruta'])
    print('  hojas          : %d' % res['hojas'])
    print('  fórmulas       : %d  (todas con valor cacheado y sin #error)'
          % res['formulas'])
    print('  celdas verdes  : %d  · vacías: %d'
          % (res['verdes'], len(res['verdes_vacias'])))
    if res['verdes_vacias']:
        print('    VACÍAS: ' + ', '.join(res['verdes_vacias'][:20]))
    print('  notas legales  : %d' % res['notas_legales'])
    print('  mapa           : %d etiquetas' % res['mapa'])
    print('  blancos OK     : %d (fórmulas cuyo resultado correcto es "")'
          % res.get('blancos_ok', 0))
    print('  inject_cache   : %s' % res['cache'])
    print('  demos          : %s' % ('OK' if ok else 'FALLAN'))
    for nombre_d, bien, detalle in demos:
        print('    [%s] %s — %s' % ('ok' if bien else 'NO', nombre_d, detalle))
    if not ok:
        raise SystemExit('Algún demo del libro no pasa.')
    return ok
