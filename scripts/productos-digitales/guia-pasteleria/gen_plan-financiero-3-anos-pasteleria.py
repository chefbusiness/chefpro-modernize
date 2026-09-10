#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
gen_plan-financiero-3-anos-pasteleria.py — libro 5 de «Cómo Montar una
Pastelería» (SPEC §2.2 fila 5; decisiones D3, D17, D21, D22, D23 y §9).

Hojas: Instrucciones · 0. Supuestos · Inversión Inicial · PyG 3 Años ·
Punto de Equilibrio · Escenarios · Personal · Tesorería 12 meses ·
Financiación · Canales y Punto Muerto.

QUÉ DECIDE ESTE LIBRO
---------------------
Si el negocio se sostiene, cuándo llega al equilibrio y QUÉ CANAL lo sostiene.
Es la FUENTE ÚNICA de las cifras del texto financiero de la guía: ninguna cifra
de euros del capítulo financiero se escribe fuera de aquí.

LIBRO HÍBRIDO (SPEC §9)
-----------------------
Calca la estructura y las fórmulas encadenadas del molde `planes-v2_0` motor 2.2
verificado en `plan-negocio-panaderia/plan-financiero-panaderia.xlsx` (9 hojas,
737 fórmulas, cero prohibidas) y le añade la hoja propia «Canales y Punto
Muerto», que absorbe el libro que la SPEC descartó (D3). **Manda la línea de
versión de la familia de GUÍAS** (1.0 · septiembre 2026), no la del motor de
planes: los gates de versión con literal de `planes-v2_0` no aplican a un
producto que no es un plan de negocio. Se declara en «Instrucciones».

DECISIONES TÉCNICAS
-------------------
* **El año 2 ES el año de crucero.** El año 1 lleva la rampa de arranque
  (`0. Supuestos` la calcula desde la tabla de la hoja de Tesorería, no la
  teclea nadie) y el año 3 crece sobre el crucero. Así la columna «Año 2»
  reproduce al céntimo `datos_ejemplo.cuenta_resultados_crucero()`, que es lo
  que citan el guion y el bonus 1, y el año 1 sigue siendo honesto.
* **Sueldo del propietario como renglón propio de los fijos** (SPEC cap. 18),
  igual que la amortización y los gastos financieros. Un negocio que sólo es
  rentable porque su dueño no cobra no es rentable.
* **El food cost del P&L es la REGLA ÚNICA del 32 %** (D23), no el food cost de
  escandallo del libro 4. La brecha entre los dos está escrita en
  «Instrucciones» con sus cuatro causas: es el supuesto conservador.
* **Margen neto de referencia 8-12 % (PS-48)** como celda de contraste, con su
  semáforo. No es un objetivo: es lo que declara un pastelero de cuarta
  generación sobre su propia casa.
* **Columna «Tipo de IVA de la línea» en toda la Inversión Inicial y en todo el
  P&L** (D17): sin ella el CAPEX sale un 21 % desviado y la tesorería paga de
  menos a los proveedores.
* **La cuota del préstamo es una anualidad algebraica** sobre un cuadro MENSUAL
  de 84 filas: la carencia de este caso son 6 meses y un cuadro anual no sabe
  expresar medio año. `PMT` está prohibida.
* Cero constantes dentro de las fórmulas; `IFERROR(...,"")` en toda división;
  «sin dato» = `""`, nunca `0`; semáforos con `ISNUMBER`; desplegables contra
  RANGO; A4 con `print_setup`; textos WinAnsi.

Salida fija: build/plan-financiero-3-anos-pasteleria.xlsx + su mapa de celdas.
Via: Claude Code
"""
import math
import os
import sys

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

AQUI = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, AQUI)

import _comun_libros_5_6 as C                                  # noqa: E402
import motor                                                   # noqa: E402
import datos_ejemplo as D                                      # noqa: E402

NOMBRE = 'plan-financiero-3-anos-pasteleria'
TITULO = 'Plan financiero a 3 años'

H_INS = 'Instrucciones'
H_SUP = '0. Supuestos'
H_INV = 'Inversión Inicial'
H_PYG = 'PyG 3 Años'
H_PEQ = 'Punto de Equilibrio'
H_ESC = 'Escenarios'
H_PER = 'Personal'
H_TES = 'Tesorería 12 meses'
H_FIN = 'Financiación'
H_CAN = 'Canales y Punto Muerto'

Q_SUP = "'" + H_SUP + "'!"
Q_INV = "'" + H_INV + "'!"
Q_PYG = "'" + H_PYG + "'!"
Q_PER = "'" + H_PER + "'!"
Q_TES = "'" + H_TES + "'!"
Q_FIN = "'" + H_FIN + "'!"
Q_CAN = "'" + H_CAN + "'!"

MESES_COL = ('B', 'C', 'D', 'E', 'F', 'G', 'H', 'I', 'J', 'K', 'L', 'M')


# ==========================================================================
# Filas fijas de cada hoja (el mapa y las fórmulas cruzadas viven de aquí)
# ==========================================================================
# --- 0. Supuestos ---------------------------------------------------------
S = {
    'tickets': 6, 'piezas': 7, 'pvp': 8, 'ticket_iva': 9, 'ticket_sin': 10,
    'dias': 11, 'factor_a1': 12, 'crec_a3': 13,
    'iva_prod': 15, 'iva_pan': 16, 'iva_gen': 17, 'iva_compras': 18,
    'peso_pan': 19, 'iva_medio': 20,
    'food_cost': 22, 'margen_obj': 23,
    'ss': 25, 'pagas': 26, 'smi': 27, 'horas_sem': 28,
    'renta': 30, 'fianza_meses': 31, 'suministros': 32, 'retribucion': 33,
    'subida': 34,
    'propios': 36, 'principal': 37, 'tipo': 38, 'plazo': 39, 'carencia': 40,
    'colchon': 41,
    'is_nueva': 43, 'is_gen': 44, 'bases_neg': 45,
    'anios_amort': 47,
    'rampa1': 49, 'rampa_meses': 50,
    'dias_cobro': 52, 'dias_pago': 53, 'extra1': 54, 'extra2': 55,
    'extra3': 56,
    'dscr_min': 58, 'dscr_obj': 59, 'holgura': 60, 'neto_suelo': 61,
    'neto_techo': 62, 'techo_personal': 63, 'techo_alquiler': 64,
    'paso_sens': 65,
}
SEC_SUP = {5: 'ACTIVIDAD', 14: 'IVA Y TIPOS', 21: 'COSTE DE VENTAS',
           24: 'PERSONAL', 29: 'LOCAL Y GASTOS FIJOS', 35: 'FINANCIACIÓN',
           42: 'FISCAL', 46: 'AMORTIZACIÓN', 48: 'ARRANQUE',
           51: 'COBROS Y PAGOS', 57: 'UMBRALES Y SUELOS DE CONTROL'}


def sup(clave):
    return Q_SUP + '$B$' + str(S[clave])


# --- Inversión Inicial ----------------------------------------------------
I_CAB = 5
I_INI = 6
I_FIN = I_INI + len(D.CAPEX) - 1            # 16
I_FONDO = I_FIN                             # el fondo de maniobra es el último
I_SUBTOT = I_FIN + 1                        # 17
I_FDM = I_FIN + 2                           # 18
I_TOTAL = I_FIN + 3                         # 19
I_CAPEX = I_FIN + 4                         # 20
I_IVA = I_FIN + 5                           # 21
I_NECES = I_FIN + 6                         # 22
I_CHK = I_FIN + 7                           # 23
I_SEC_AM = I_FIN + 9                        # 25
I_BASE = I_FIN + 10                         # 26
I_AMORT = I_FIN + 11                        # 27
I_NOTA = I_FIN + 13                         # 29
_FILA_CAPEX = dict((i, I_INI + i) for i in range(len(D.CAPEX)))
I_FIANZA = _FILA_CAPEX[7]
I_PACK = _FILA_CAPEX[8]

# --- PyG 3 Años -----------------------------------------------------------
P = {'cab': 5, 'sec_ing': 6, 'tickets': 7, 'ticket': 8, 'dias': 9,
     'actividad': 10, 'ingresos': 11, 'v_past': 12, 'v_pan': 13,
     'sec_var': 14, 'coste_ventas': 15, 'tot_var': 16, 'margen': 17,
     'sec_fijos': 18, 'personal': 19, 'propietario': 20, 'alquiler': 21,
     'suministros': 22, 'seguros': 23, 'gestoria': 24, 'software': 25,
     'telefonia': 26, 'limpieza': 27, 'mantenimiento': 28, 'publicidad': 29,
     'otros': 30, 'amortizacion': 31, 'financieros': 32, 'tot_fijos': 33,
     'iva_var': 34, 'iva_fijos': 35, 'rai': 36, 'bases_ini': 37,
     'base_imp': 38, 'ejercicios': 39, 'tipo_is': 40, 'is': 41,
     'bases_fin': 42, 'neto': 43,
     'sec_ratios': 44, 'cab_ratios': 45, 'r_margen': 46, 'r_coste': 47,
     'r_personal': 48, 'r_alquiler': 49, 'r_neto': 50, 'r_banda': 51,
     'r_equilibrio': 52, 'nota': 54}
FIJOS_VERDES = [
    ('seguros', 'Seguros (responsabilidad civil y continente)', 1440.0, 0.0,
     'Supuesto. El de RC no está verificado como obligatorio: depende de tu '
     'CCAA. Las operaciones de seguro están exentas de IVA (art. 20.Uno.16 de '
     'la Ley 37/1992), y por eso su tipo va a cero.'),
    ('gestoria', 'Gestoría y asesoría', 2640.0, 0.21, 'Supuesto.'),
    ('software', 'Software, TPV y pasarela de cobro', 720.0, 0.21,
     'Supuesto. Con Verifactu en el horizonte esta cuota deja de ser opcional.'),
    ('telefonia', 'Telefonía e internet', 660.0, 0.21, 'Supuesto.'),
    ('limpieza', 'Limpieza y consumibles', 2160.0, 0.21, 'Supuesto.'),
    ('mantenimiento', 'Mantenimiento de equipos', 1800.0, 0.21,
     'Supuesto. Con horno de gas, súmale la inspección periódica obligatoria '
     'cada 5 años, que se te repercute.'),
    ('publicidad', 'Publicidad y redes', 2400.0, 0.21, 'Supuesto.'),
    ('otros', 'Otros gastos de estructura', 1800.0, 0.21, 'Supuesto.'),
]

# --- Punto de Equilibrio --------------------------------------------------
E = {'cab': 5, 'sec_base': 6, 'fijos_anio': 7, 'fijos_mes': 8, 'ticket': 9,
     'cv_ticket': 10, 'mc_ticket': 11, 'dias': 12, 'principal': 13,
     'sec_cont': 14, 'tk_anio': 15, 'tk_dia': 16, 'ing_anio': 17,
     'ing_mes': 18,
     'sec_caja': 19, 'fijos_caja': 20, 'tkc_anio': 21, 'tkc_dia': 22,
     'ingc_anio': 23,
     'sec_contraste': 24, 'tk_previstos': 25, 'holgura': 26, 'semaforo': 27,
     'nota': 28,
     'sec_sens': 30, 'sens_cab': 31, 'sens_ini': 32, 'sens_fin': 36,
     'sens_nota': 38}

# --- Escenarios -----------------------------------------------------------
X = {'cab': 5, 'tickets': 6, 'ticket': 7, 'dias': 8, 'ingresos': 9,
     'coste': 10, 'tot_var': 11, 'margen': 12, 'fijos': 13, 'rai': 14,
     'is': 15, 'neto': 16, 'margen_neto': 17, 'tk_equilibrio': 18,
     'sec_exige': 19, 'personal_ventas': 20, 'ventas_jornada': 21,
     'saldo_est': 22, 'saldo_real': 23, 'nota': 25}

# --- Personal -------------------------------------------------------------
R_CAB = 5
R_INI = 6
R_FIN = R_INI + len(D.PLANTILLA) - 1        # 10
R_TOT = R_FIN + 1                           # 11
R_SEC_CONV = R_TOT + 2                      # 13
R_CONV_CAB = R_SEC_CONV + 1                 # 14
R_CONV_INI = R_CONV_CAB + 1                 # 15
R_CONV_FIN = R_CONV_INI + len(D.CONVENIO) - 1   # 20
R_SEC_SMI = R_CONV_FIN + 2                  # 22
R_SUELO = R_SEC_SMI + 1                     # 23
R_SMI = R_SEC_SMI + 2                       # 24
R_VEREDICTO = R_SEC_SMI + 3                 # 25
R_SEC_HORAS = R_SEC_SMI + 5                 # 27
R_HORAS = R_SEC_HORAS + 1                   # 28
R_JORNADAS = R_SEC_HORAS + 2                # 29
R_AVISOS = R_SEC_HORAS + 3                  # 30
R_NOTA = R_SEC_HORAS + 5                    # 32

# --- Tesorería 12 meses ---------------------------------------------------
T = {'cab': 5, 'estacionalidad': 6, 'rampa': 7, 'reparto': 8,
     'sec_cobros': 9, 'ventas': 10,
     'sec_pagos': 11, 'compras': 12, 'fijos': 13, 'nominas': 14,
     'intereses': 15, 'principal': 16,
     'iva_rep': 17, 'iva_sop': 18, 'iva_arr': 19, 'iva_liq': 20, 'iva_pago': 21,
     'flujo': 22, 'saldo': 23, 'saldo_min': 24, 'mes_fondo': 25,
     'sec_retorno': 27, 'fcf1': 28, 'fcf2': 29, 'fcf3': 30, 'inversion': 31,
     'payback': 32, 'nota': 34}

# --- Financiación ---------------------------------------------------------
F = {'sec_origen': 5, 'propios': 6, 'prestamo': 7, 'otras': 8, 'tot_origen': 9,
     'sec_usos': 10, 'necesidad': 11, 'diferencia': 12, 'dif_pct': 13,
     'ajuste': 14,
     'sec_cond': 15, 'principal': 16, 'tipo': 17, 'tipo_mes': 18,
     'plazo': 19, 'carencia': 20, 'meses_am': 21, 'cuota': 22, 'nota_cuota': 23,
     'cab': 25}
F_MES_INI = 26
F_MESES = 84
F_MES_FIN = F_MES_INI + F_MESES - 1         # 109
F_TOL = F_MES_FIN + 1                       # 110
F_PENDIENTE = F_MES_FIN + 2                 # 111
F_CIERRA = F_MES_FIN + 3                    # 112
F_SEC_ANIO = F_MES_FIN + 5                  # 114
F_ANIO_CAB = F_MES_FIN + 6                  # 115
F_ANIO_INI = F_MES_FIN + 7                  # 116
F_ANIOS = 7
F_ANIO_FIN = F_ANIO_INI + F_ANIOS - 1       # 122
F_DSCR_MIN = F_ANIO_FIN + 1                 # 123
F_DSCR_SEM = F_ANIO_FIN + 2                 # 124
F_NOTA = F_ANIO_FIN + 4                     # 126

# --- Canales y Punto Muerto -----------------------------------------------
K_CAB = 5
K_INI = 6
K_FIN = K_INI + len(D.CANALES) - 1          # 9
K_TOT = K_FIN + 1                           # 10
K_B2B = K_INI + 2                           # fila del canal B2B (índice 2)
K_SEC_PM = K_TOT + 2                        # 12
K_FIJOS = K_SEC_PM + 1                      # 13
K_FIJOS_MES = K_SEC_PM + 2                  # 14
K_MC = K_SEC_PM + 3                         # 15
K_REGLA = K_SEC_PM + 4                      # 16
K_DIF_MC = K_SEC_PM + 5                     # 17
K_PM_CON = K_SEC_PM + 6                     # 18
K_PM_SIN = K_SEC_PM + 7                     # 19
K_PM_DIF = K_SEC_PM + 8                     # 20
K_SOSTIENE = K_SEC_PM + 9                   # 21
K_APORTE = K_SEC_PM + 10                    # 22
K_SEC_ART3 = K_SEC_PM + 12                  # 24
K_PESO_B2B = K_SEC_ART3 + 1                 # 25
K_UMBRAL = K_SEC_ART3 + 2                   # 26
K_VEREDICTO = K_SEC_ART3 + 3                # 27
K_COBRO = K_SEC_ART3 + 5                    # 29
K_NOTA = K_SEC_ART3 + 7                     # 31


def ie(expr):
    return motor.iferror(expr)


# ==========================================================================
# Hoja «Instrucciones»
# ==========================================================================
PASOS = [
    '1. Hoja «0. Supuestos»: es el cuadro de mandos. Aquí se teclean las TASAS '
    'y los DRIVERS —tickets al día, ticket medio, días de apertura, tipos de '
    'IVA, food cost objetivo, condiciones del préstamo, umbrales de control—, y '
    'el resto del libro se recalcula solo. Las partidas de gasto se teclean en '
    'su hoja, no aquí.',
    '2. Hoja «Inversión Inicial»: cuánto cuesta abrir, partida a partida, con '
    'su columna «¿Lleva IVA?» y su TIPO por línea. Sin esa columna el CAPEX '
    'sale hasta un 21 % desviado, que es el error más caro de todo el proceso. '
    'Abajo salen el IVA que hay que ADELANTAR y la necesidad total de caja.',
    '3. Hoja «PyG 3 Años»: la cuenta de resultados. El AÑO 2 es el año de '
    'crucero; el año 1 lleva la rampa de arranque y el año 3 crece sobre el '
    'crucero. El sueldo del propietario, la amortización y los intereses van '
    'DENTRO de los costes fijos: si no, la rentabilidad que publica un plan de '
    'negocio no significa nada.',
    '4. Hoja «Punto de Equilibrio»: cuántos tickets al día hay que hacer para '
    'no perder dinero, en dos versiones. La CONTABLE incluye la amortización; '
    'la de CAJA le quita la amortización (que no se paga) y le suma la '
    'devolución de principal (que sí). La segunda es la que decide si llegas a '
    'fin de mes.',
    '5. Hoja «Escenarios»: el mismo modelo con tres juegos de tickets, ticket '
    'medio y días de apertura. La columna «Realista» lee el año de crucero del '
    'P&L, así que no puede desviarse de él.',
    '6. Hoja «Personal»: cinco personas y tres jornadas y media, con el bruto '
    'del convenio en celda verde y la Seguridad Social a cargo de la empresa '
    'DENTRO. Del bruto al coste empresa hay un tercio más: es el salto que más '
    'sorprende.',
    '7. Hoja «Tesorería 12 meses»: el P&L dice si el negocio gana dinero; esta '
    'hoja dice si le queda caja para llegar a fin de mes. Trae la '
    'estacionalidad mes a mes, la rampa de arranque, el desfase de cobros y '
    'pagos y la liquidación trimestral del IVA.',
    '8. Hoja «Financiación»: de dónde sale el dinero y cuánto cuesta '
    'devolverlo. El cuadro de amortización es MENSUAL, con 84 filas, porque la '
    'carencia de este caso son seis meses y un cuadro anual no sabe expresar '
    'medio año. Abajo, el DSCR año a año: es el número que decide la '
    'operación bancaria.',
    '9. Hoja «Canales y Punto Muerto»: mostrador, encargos, B2B y online, cada '
    'uno con su margen, su coste de servir, su comisión y sus días de cobro. '
    'Dice qué canal sostiene el negocio y qué pasa con el punto muerto si '
    'quitas el B2B. Y avisa de cuándo ese B2B te mete en el art. 3 del '
    'RD 1021/2022.',
]

NOTAS_LIBRO = [
    'ESTE LIBRO ES LA FUENTE ÚNICA DE LAS CIFRAS FINANCIERAS DE LA GUÍA. Si un '
    'euro del texto no sale de una celda de aquí, es que está inventado. '
    'Cámbialo aquí y el capítulo entero deja de cuadrar contigo: eso es lo que '
    'tiene que pasar.',
    'LIBRO HÍBRIDO. La estructura y las fórmulas encadenadas son las del molde '
    'de planes de negocio de la casa (nueve hojas, con tesorería, IVA y '
    'financiación de verdad), y encima va la hoja «Canales y Punto Muerto», que '
    'es propia de esta guía. La versión que manda es la de la familia de guías, '
    'la que está al pie de esta hoja, no la del motor de planes.',
    'EL FOOD COST DEL P&L NO ES EL DE TU ESCANDALLO, Y NO ES UN ERROR. El '
    'escandallo del libro «carta-de-apertura-y-escandallo» mide la materia '
    'prima de la pieza; la regla del 32 % que proyecta este P&L incluye además '
    'cuatro cosas que el escandallo de la pieza deja fuera: el producto de '
    'terceros que revende el despacho (bebida, café, helado, bombonería '
    'comprada), la merma real del producto fresco, el packaging y la subida de '
    'la materia prima. Por eso el plan se proyecta con el 32 % y no con el '
    'escandallo: es el supuesto CONSERVADOR. La diferencia entre los dos '
    'números es tu colchón.',
    'EL SUELDO DEL PROPIETARIO ES UN COSTE. Va como renglón propio de los '
    'fijos, con la cuota de autónomos dentro, y no está en la plantilla de '
    'cinco personas. Un negocio que sólo es rentable porque su dueño no cobra '
    'no es rentable: es un empleo mal pagado con riesgo patrimonial.',
    'EL 8-12 % DE MARGEN NETO NO ES UN OBJETIVO: es lo que declara sobre su '
    'propia casa un pastelero de cuarta generación. Está en el libro como celda '
    'de contraste con su semáforo, para que veas si tu proyección se sale por '
    'arriba (y entonces desconfía de ella) o por abajo.',
    'TODAS LAS CIFRAS VAN SIN IVA salvo donde la hoja dice lo contrario. La '
    'tesorería es caja: ahí el IVA entra y sale, y la liquidación trimestral lo '
    'devuelve a su sitio. Las columnas de los años 2 y 3 están en euros del año '
    '1 mientras la «Subida anual de los costes fijos» de Supuestos valga cero.',
]

FRONTERA = [
    ('kit-tareas-pasteleria (12 €)',
     'La operación del obrador que ya tienes: plan de producción semanal, '
     'encargos, alérgenos de vitrina, temperaturas.',
     'Este libro NO planifica producción ni controla encargos. Toma su '
     'resultado en euros y lo proyecta a tres años.'),
    ('kit-escandallos (12 €) y el libro «carta-de-apertura-y-escandallo»',
     'El coste real de cada pieza, con su mano de obra y su margen.',
     'Aquí el coste de ventas entra como UN porcentaje sobre las ventas: la '
     'regla del 32 %. Cuando midas el tuyo, cámbialo en «0. Supuestos».'),
    ('plan-negocio-panaderia (35 €)',
     'El dossier completo que se le entrega al banco, con memoria escrita.',
     'Este libro es el motor de números de una GUÍA de apertura. Si vas a '
     'pedir financiación de verdad, el dossier del banco es aquel producto.'),
    ('El libro «calculadora-capex-pasteleria» de esta misma guía',
     'El CAPEX partida a partida, la variante del formato y la comparación '
     'entre traspaso y obra nueva.',
     'Aquí sólo entran los TOTALES por bloque. Las dos celdas de equipamiento '
     'de la hoja de Inversión vienen sembradas y son verdes: sustitúyelas por '
     'el total de aquel libro cuando lo tengas cerrado.'),
]


def hoja_instrucciones(wb):
    ws = wb.create_sheet(H_INS, 0)
    C.anchos(ws, {'A': 44, 'B': 44, 'C': 44})
    motor.val(ws, 'A1', TITULO)
    ws['A1'].font = Font(bold=True, size=16, color=C.ORO)
    ws.row_dimensions[1].height = 26
    motor.val(ws, 'A2', C.SUBTITULO)
    ws['A2'].font = Font(size=9, color=C.GRIS_TXT)
    motor.val(ws, 'A3', 'Para qué sirve: saber si el negocio se sostiene, '
                        'cuándo llega al equilibrio y qué canal lo sostiene.')
    ws['A3'].font = Font(italic=True, size=9)

    fila = 5
    C.seccion(ws, 'A%d' % fila, 'Instrucciones de uso')
    fila += 1
    for paso in PASOS:
        C.parrafo(ws, fila, paso, 'A', 'C', alto=54)
        fila += 1
    fila += 1
    motor.val(ws, 'A%d' % fila, C.NOTA_VERDES)
    ws['A%d' % fila].fill = PatternFill('solid', fgColor=motor.VERDE)
    fila += 2

    C.seccion(ws, 'A%d' % fila, 'Lo que conviene saber antes de empezar')
    fila += 1
    for texto in NOTAS_LIBRO:
        C.parrafo(ws, fila, texto, 'A', 'C', alto=68)
        fila += 1
    fila += 1

    C.seccion(ws, 'A%d' % fila,
              'QUÉ HACE ESTE LIBRO Y QUÉ HACE OTRO (para no teclear dos veces)')
    fila += 1
    C.cabecera(ws, fila, [('A', 'Producto o libro'), ('B', 'Qué aporta'),
                          ('C', 'Qué NO hace este libro')], altura=26)
    fila += 1
    for producto, aporta, no_hace in FRONTERA:
        motor.val(ws, 'A%d' % fila, producto, bold=True, wrap=True)
        motor.val(ws, 'B%d' % fila, aporta, wrap=True)
        motor.val(ws, 'C%d' % fila, no_hace, wrap=True)
        ws.row_dimensions[fila].height = 62
        fila += 1
    C.parrafo(ws, fila,
              'En todo este paquete no hay ni una fórmula que apunte a otro '
              'fichero. Lo que viene de otro libro viene COPIADO y declarado: '
              'un enlace entre ficheros se rompe en cuanto alguien mueve una '
              'carpeta, y entonces el libro miente sin avisar.', 'A', 'C',
              alto=42)
    fila += 2

    C.seccion(ws, 'A%d' % fila, 'Cada cuánto se usa este libro')
    fila += 1
    C.parrafo(ws, fila,
              'Cadencia: «0. Supuestos» e «Inversión Inicial», MIENTRAS BUSCAS '
              'LOCAL y pides presupuestos, tantas veces como haga falta. El '
              'P&L, el punto de equilibrio y los escenarios, UNA VEZ ANTES DE '
              'FIRMAR y otra antes de sentarte con el banco. «Tesorería 12 '
              'meses», CADA MES durante el primer año, sustituyendo la '
              'previsión por lo que pasó de verdad. «Canales y Punto Muerto», '
              'cada vez que te plantees abrir o cerrar un canal.', 'A', 'C',
              alto=68)
    fila += 2

    C.parrafo(ws, fila, C.NOTA_DESPROTEGER, 'A', 'C', alto=28)
    fila += 2
    C.pie(ws, fila, 'A', 'C')
    C.pagina(ws, apaisado=False, area='A1:C%d' % (fila + 1))
    return ws


# ==========================================================================
# Hoja «0. Supuestos»
# ==========================================================================
def hoja_supuestos(wb):
    ws = wb.create_sheet(H_SUP)
    C.anchos(ws, {'A': 54, 'B': 16, 'C': 96})
    C.encabezar(ws, 'Supuestos: las tasas y los drivers del modelo',
                'Cambia las celdas VERDES y el resto del libro se recalcula '
                'solo. Las partidas de gasto se teclean en su hoja, no aquí. '
                'Cada supuesto lleva al lado de dónde sale y qué pasa si lo '
                'mueves.', col_fin='C')
    for f, texto in SEC_SUP.items():
        C.seccion(ws, 'A%d' % f, texto)
        motor.val(ws, 'B%d' % f, '')
        motor.val(ws, 'C%d' % f, '')
        ws['B%d' % f].fill = PatternFill('solid', fgColor=C.CABECERA)
        ws['C%d' % f].fill = PatternFill('solid', fgColor=C.CABECERA)

    def ent(clave, etiqueta, valor, fmt, nota_txt, id_pa=None):
        fila = S[clave]
        motor.val(ws, 'A%d' % fila, etiqueta)
        C.entrada(ws, 'B%d' % fila, valor, fmt=fmt, etiqueta=etiqueta)
        C.nota(ws, 'C%d' % fila, nota_txt)
        if id_pa:
            C.nota_celda(ws, 'B%d' % fila, id_pa)
        return 'B%d' % fila

    def fx(clave, etiqueta, formula, fmt, nota_txt):
        fila = S[clave]
        motor.val(ws, 'A%d' % fila, etiqueta)
        motor.f(ws, 'B%d' % fila, formula, fmt=fmt)
        C.nota(ws, 'C%d' % fila, nota_txt)
        return 'B%d' % fila

    # --- ACTIVIDAD --------------------------------------------------------
    ent('tickets', 'Tickets al día en velocidad de crucero',
        D.P('tickets_dia_crucero'), C.ENT,
        'SUPUESTO. No existe dato público de tráfico de pastelería en España, '
        'así que este número no es un benchmark: es el que hace que el modelo '
        'tenga qué calcular. Cuéntalos tú en la puerta de tres pastelerías '
        'parecidas a la que quieres abrir, un sábado y un martes.')
    ent('piezas', 'Piezas por ticket', D.P('piezas_por_ticket'), C.DEC1,
        'SUPUESTO. De aquí sale el ticket medio, que NO se teclea: se calcula. '
        'Publicar un «ticket medio de pastelería» sería inventarse un dato que '
        'no existe.')
    ent('pvp', 'PVP medio ponderado de la carta, CON IVA (€/pieza)',
        round(D.pvp_medio_ponderado(), 4), C.EUR,
        'Sale de las 30 referencias del libro «carta-de-apertura-y-escandallo», '
        'ponderadas por su mix de ventas. Si cambias la carta allí, trae aquí '
        'el nuevo PVP medio: es COPIA declarada, no un vínculo entre ficheros.')
    fx('ticket_iva', 'Ticket medio CON IVA (€)',
       ie('$B$%d*$B$%d' % (S['piezas'], S['pvp'])), C.EUR,
       'Piezas por ticket x PVP medio ponderado. Es lo que ve el cliente en el '
       'datáfono.')
    fx('ticket_sin', 'Ticket medio SIN IVA (€)',
       ie('$B$%d/(1+$B$%d)' % (S['ticket_iva'], S['iva_medio'])), C.EUR,
       'Es el que entra en el P&L: la cuenta de resultados va siempre sin IVA. '
       'Se divide por el IVA MEDIO de la carta, no por un tipo suelto, porque '
       'los panes van al 4 % y el resto al 10 %.')
    ent('dias', 'Días de apertura al año', D.NEGOCIO['dias_apertura_anio'],
        C.ENT,
        'SUPUESTO: 6 días por semana x 52 semanas = 312, menos 12 días de '
        'cierre en agosto. Los festivos NO se descuentan: la pastelería tiene '
        'libertad horaria por ley estatal cuando es la actividad principal. El '
        'MISMO dato lo usan el P&L, el punto de equilibrio y los escenarios.')
    fx('factor_a1', 'Actividad del año 1 sobre la de crucero (%)',
       ie("SUMPRODUCT(%sB%d:M%d,%sB%d:M%d)"
          % (Q_TES, T['estacionalidad'], T['estacionalidad'],
             Q_TES, T['rampa'], T['rampa'])), C.PCT,
       'NO se teclea: sale de la estacionalidad y de la rampa de la hoja de '
       'Tesorería. Un local que abre no factura desde el primer día lo que '
       'factura a los seis meses, y proyectar el año 1 a velocidad de crucero '
       'es el error que más planes de negocio tumba.')
    ent('crec_a3', 'Crecimiento de volumen del año 3 (%)', 0.04, C.PCT,
        'SUPUESTO. El año 2 ya es el de crucero: el año 3 sólo crece lo que dé '
        'la clientela fija, no un salto de negocio. Un 4 % es prudente; si '
        'pones más, tendrás que decir de dónde sale.')

    # --- IVA --------------------------------------------------------------
    ent('iva_prod', 'IVA de pastelería, bollería y confitería (%)',
        D.P('iva_producto'), C.PCT,
        'El 10 % sale de la regla general de alimentos del art. 91.Uno.1.1.o '
        'de la Ley 37/1992: no porque un precepto los nombre, sino porque no '
        'están en la lista cerrada del 4 % ni excluidos del 10 %.', 'PA-36')
    ent('iva_pan', 'IVA del pan (%)', D.P('iva_pan'), C.PCT,
        'TODOS los productos del RD 308/2019 (pan común, pan especial y '
        'semielaborados, con o sin gluten) van al 4 %. El cruasán sigue al '
        '10 %: el RD 308/2019 sólo regula el pan y no menciona la bollería.',
        'PA-36b')
    ent('iva_gen', 'IVA general (%)', D.P('iva_general'), C.PCT,
        'Tipo general del art. 90.Uno de la Ley 37/1992. Es el que soportas en '
        'obra, maquinaria, packaging y servicios.')
    ent('iva_compras', 'IVA medio de las compras de materia prima (%)',
        D.P('iva_producto'), C.PCT,
        'SUPUESTO de trabajo: harinas y azúcares se mueven entre el 4 % y el '
        '10 %, chocolate y mantequilla al 10 %, y el packaging al 21 %. Sólo '
        'afecta a la TESORERÍA (lo que adelantas al proveedor), nunca al '
        'resultado. Píde a tu distribuidor una factura de ejemplo y afínalo.')
    ent('peso_pan', 'Peso del pan sobre las ventas sin IVA (%)',
        round(0.20127639, 6), C.PCT2,
        'Sale de la familia «panes de acompañamiento» de la carta del libro 4, '
        'ponderada por su mix. Existe porque el pan tributa al 4 % y el resto '
        'al 10 %: sin este reparto, la caja y las ventas netas no cuadran.')
    fx('iva_medio', 'IVA medio ponderado de la carta (%)',
       ie('$B$%d*$B$%d+(1-$B$%d)*$B$%d'
          % (S['peso_pan'], S['iva_pan'], S['peso_pan'], S['iva_prod'])),
       C.PCT2,
       'No es un tipo legal: es el que hace falta para pasar de caja a ventas '
       'netas cuando vendes a dos tipos distintos.')

    # --- COSTE DE VENTAS --------------------------------------------------
    ent('food_cost', 'Coste de ventas sobre las ventas (food cost, %)',
        D.P('food_cost_objetivo'), C.PCT,
        'REGLA ÚNICA de la casa: margen bruto del 65-70 % equivale a un food '
        'cost del 30-35 %, y son la misma regla dicha dos veces. El 32 % es el '
        'punto medio. Incluye lo que el escandallo de la pieza deja fuera: '
        'producto de terceros que revende el despacho, merma real, packaging y '
        'subida de materia prima. Por eso es MAYOR que tu escandallo, y por eso '
        'es el supuesto conservador.')
    fx('margen_obj', 'Margen bruto objetivo (%)',
       ie('1-$B$%d' % S['food_cost']), C.PCT,
       'Se DERIVA del food cost, no se guarda aparte: si se guardaran los dos, '
       'antes o después dirían cosas distintas.')

    # --- PERSONAL ---------------------------------------------------------
    motor.escribir_parametro(ws, S['ss'], 'A', 'B', 'ss_empresa', col_nota='C')
    C.VERDES.append((ws.title, 'B%d' % S['ss'],
                     motor.PARAMETROS['ss_empresa']['etiqueta'],
                     motor.PARAMETROS['ss_empresa']['valor']))
    ent('pagas', 'Número de pagas del convenio', D.P('pagas_convenio'),
        C.ENT,
        'El convenio de pastelería de Madrid paga 15, no 14, y eso cambia el '
        'coste mes a mes y el reparto de la tesorería. Busca el tuyo: el '
        'convenio es provincial o autonómico.', 'PA-33')
    motor.escribir_parametro(ws, S['smi'], 'A', 'B', 'smi_anual', col_nota='C')
    C.VERDES.append((ws.title, 'B%d' % S['smi'],
                     motor.PARAMETROS['smi_anual']['etiqueta'],
                     motor.PARAMETROS['smi_anual']['valor']))
    C.nota_celda(ws, 'B%d' % S['smi'], 'PA-32')
    ent('horas_sem', 'Jornada completa (horas/semana)',
        D.P('horas_semana_jornada_completa'), C.ENT,
        'SUPUESTO: la jornada anual del convenio de Madrid no se ha '
        'verificado. Comprueba la de tu convenio antes de dimensionar turnos.')

    # --- LOCAL Y FIJOS ----------------------------------------------------
    ent('renta', 'Alquiler mensual del local (€)', D.NEGOCIO['renta_mensual'],
        C.EUR,
        'SUPUESTO para una ciudad media. Los cinco traspasos de Barcelona del '
        'research se mueven entre 850 y 1.300 €/mes para 63-85 m2; aquí se '
        'toma la parte baja porque Barcelona no es una ciudad media.')
    ent('fianza_meses', 'Fianza del alquiler (meses de renta)',
        D.NEGOCIO['meses_fianza'], C.ENT,
        'SUPUESTO. La fianza no es un gasto: es un depósito que se recupera. '
        'Va en la inversión porque hay que tenerla el día de la firma.')
    ent('suministros', 'Suministros mensuales de luz, gas y agua (€)',
        1450.0, C.EUR,
        'SUPUESTO. Con 62 kW instalados y frío 24 horas es la segunda partida '
        'del negocio. No se estima: se le pide la simulación a la '
        'comercializadora con la potencia del proyecto eléctrico.')
    ent('retribucion', 'Retribución mensual del propietario (€)', 2400.0,
        C.EUR,
        'SUPUESTO, y renglón propio de los fijos con la cuota de autónomos '
        'dentro. Si el dueño está en el obrador, su trabajo tiene un coste '
        'aunque no se lo pague a fin de mes.')
    ent('subida', 'Subida anual de los costes fijos (%)', 0.0, C.PCT,
        'A CERO: las tres columnas del P&L están en euros del año 1 (términos '
        'reales). Súbela si quieres proyectar en euros corrientes, pero '
        'entonces sube también el ticket medio o estarás proyectando una '
        'pérdida de margen que no has decidido.')

    # --- FINANCIACIÓN -----------------------------------------------------
    # Se redondea HACIA ARRIBA al céntimo: sembrarlo con un redondeo normal
    # dejaba la diferencia entre origen y usos en -0,004 €, lo bastante para
    # encender el semáforo rojo de «el plan no está financiado» en un libro
    # recién abierto.
    propios = (math.ceil((_necesidad_estimada()
                          - D.FINANCIACION['principal']) * 100.0) / 100.0)
    ent('propios', 'Recursos propios aportados (€)', propios, C.EUR,
        'SUPUESTO sembrado para que el plan quede financiado: es la necesidad '
        'total de caja menos el préstamo. Míralo dos veces, porque es mucho '
        'dinero: si no lo tienes, o baja la inversión (mira la hoja «Traspaso '
        'vs Obra Nueva» del libro de CAPEX) o sube el préstamo, y entonces '
        'vigila el DSCR de la hoja de Financiación.')
    ent('principal', 'Préstamo bancario solicitado (€)',
        D.FINANCIACION['principal'], C.EUR,
        'SUPUESTO. Las condiciones las pone tu banco y dependen de la '
        'garantía. Pide oferta a dos entidades y a una línea ICO antes de '
        'fijarlo.')
    ent('tipo', 'Tipo de interés nominal anual (%)',
        D.FINANCIACION['tipo_nominal'], C.PCT2, 'SUPUESTO.')
    ent('plazo', 'Plazo total del préstamo (meses)',
        D.FINANCIACION['plazo_meses'], C.ENT,
        'SUPUESTO: 7 años, carencia incluida. El cuadro de amortización de la '
        'hoja de Financiación tiene 84 filas; si alargas el plazo por encima '
        'de 84 meses, tendrás que añadirlas.')
    ent('carencia', 'Carencia de principal (meses)',
        D.FINANCIACION['carencia_meses'], C.ENT,
        'SUPUESTO, y es lo PRIMERO que hay que negociar: la carencia es lo que '
        'hace que el préstamo no se coma la caja justo en la rampa de '
        'arranque. Durante la carencia sólo se pagan intereses.')
    ent('colchon', 'Fondo de maniobra (meses de costes fijos)',
        D.P('meses_colchon_fondo_maniobra'), C.ENT,
        'SUPUESTO. Meses de gastos fijos que hay que tener en caja el día que '
        'abres. Es una partida de la inversión, no una propina.')

    # --- FISCAL -----------------------------------------------------------
    ent('is_nueva', 'Impuesto de Sociedades, entidad de nueva creación (%)',
        0.15, C.PCT,
        'Art. 29.1 de la Ley del Impuesto sobre Sociedades: se aplica al '
        'PRIMER ejercicio con base imponible positiva y al siguiente. Si eres '
        'autónomo y no sociedad, esta línea no te aplica: tributas en IRPF.')
    ent('is_gen', 'Impuesto de Sociedades, tipo general (%)', 0.25, C.PCT,
        'A partir del tercer ejercicio con base positiva.')
    ent('bases_neg', 'Bases negativas de ejercicios anteriores (€)', 0.0,
        C.EUR,
        'Pérdidas pendientes de compensar al empezar. En una apertura son '
        'cero: se rellena si arrancas dentro de una sociedad que ya existía.')

    # --- AMORTIZACIÓN -----------------------------------------------------
    ent('anios_amort', 'Años de amortización del inmovilizado',
        D.ANIOS_AMORTIZACION, C.ENT,
        'SUPUESTO, y en la práctica obra y maquinaria tienen coeficientes '
        'distintos. Sin amortización la rentabilidad publicada es mentira: los '
        'hornos se gastan.')

    # --- ARRANQUE ---------------------------------------------------------
    ent('rampa1', 'Actividad del mes 1 sobre la de crucero (%)',
        D.RAMPA['mes1'], C.PCT,
        'SUPUESTO. Con el 100 % eliminas la rampa y el año 1 se proyecta como '
        'si abrieras a pleno rendimiento, que es justo lo que no pasa.')
    ent('rampa_meses', 'Meses hasta alcanzar el régimen de crucero',
        D.RAMPA['meses'], C.ENT,
        'SUPUESTO. La rampa sube en línea recta desde el porcentaje de arriba '
        'hasta el 100 % en este número de meses.')

    # --- COBROS Y PAGOS ---------------------------------------------------
    fx('dias_cobro', 'Días medios de cobro',
       ie('SUMPRODUCT(%sB%d:B%d,%sF%d:F%d)'
          % (Q_CAN, K_INI, K_FIN, Q_CAN, K_INI, K_FIN)), C.DEC1,
       'NO se teclea: sale de la hoja «Canales y Punto Muerto», ponderando los '
       'días de cobro de cada canal por su peso en las ventas. El mostrador '
       'cobra al contado; el B2B, a 45 días, y ahí financias tú.')
    ent('dias_pago', 'Días medios de pago a proveedor',
        30, C.ENT,
        'SUPUESTO. Es lo que te financian tus proveedores, y lo primero que se '
        'pierde si te retrasas en un pago.')
    ent('extra1', 'Mes de la primera paga extra (1-12)', 7, C.ENT,
        'SUPUESTO. Con 15 pagas hay TRES meses en los que la nómina sale '
        'doble; escribe aquí en cuáles cae en tu convenio.')
    ent('extra2', 'Mes de la segunda paga extra (1-12)', 12, C.ENT,
        'SUPUESTO.')
    ent('extra3', 'Mes de la tercera paga extra (1-12)', 3, C.ENT,
        'SUPUESTO. Si tu convenio paga 14, pon aquí un número fuera de 1-12 '
        '(por ejemplo 0) y baja a 14 las pagas de arriba.')

    # --- UMBRALES ---------------------------------------------------------
    ent('dscr_min', 'DSCR mínimo aceptable', 1.0, C.DEC,
        'Por debajo de 1 el negocio no genera lo suficiente para pagar el '
        'préstamo. Pon el que te exija tu banco en el contrato.')
    ent('dscr_obj', 'DSCR objetivo (verde)', 1.25, C.DEC,
        'Lo que suele pedir una entidad para dar el préstamo sin garantías '
        'adicionales.')
    ent('holgura', 'Holgura mínima sobre el punto de equilibrio (%)', 0.15,
        C.PCT, 'Cuánto puedes caer sobre lo previsto antes de entrar en '
        'pérdidas.')
    ent('neto_suelo', 'Margen neto de referencia, suelo (%)', 0.08, C.PCT,
        'Rentabilidad neta de una pastelería «yendo bien» declarada por el '
        'gerente de una casa de cuarta generación: 8-12 %. No es un objetivo '
        'ni un benchmark del sector: es un testimonio, y está aquí para que '
        'contrastes tu proyección.')
    ent('neto_techo', 'Margen neto de referencia, techo (%)', 0.12, C.PCT,
        'Si tu proyección sale por encima del 12 %, desconfía de ella antes de '
        'enseñarla: casi siempre falta un coste.')
    ent('techo_personal', 'Techo de coste de personal sobre ventas (%)', 0.35,
        C.PCT,
        'SUPUESTO de control. En un obrador con despacho el personal es la '
        'primera partida; por encima de este techo, o subes precios o '
        'redimensionas turnos.')
    ent('techo_alquiler', 'Techo de alquiler sobre ventas (%)', 0.10, C.PCT,
        'SUPUESTO de control. Un alquiler por encima del 10 % de las ventas se '
        'come el margen de todo lo demás.')
    ent('paso_sens', 'Paso de la sensibilidad del coste variable (€)', 0.20,
        C.EUR,
        'Cuánto sube y baja el coste variable por ticket en la tabla de '
        'sensibilidad del punto de equilibrio.')

    ws.freeze_panes = 'A6'          # hallazgo B12 (2026-09-10)
    C.pagina(ws, apaisado=False, titulos='$5:$5', area='A1:C%d' % S['paso_sens'])
    return ws


# ==========================================================================
# Hoja «Inversión Inicial»
# ==========================================================================
DEFECTO_CAPEX = {
    1: round(D.equipamiento_bloque_sin_iva('Equipamiento de obrador'), 2),
    2: round(D.equipamiento_bloque_sin_iva('Tienda y vitrina'), 2),
}
NOTA_EQUIPO = ('Total del bloque, sembrado desde la dotación tipo del libro '
               '«checklist-equipamiento-y-proveedores», TODO llevado a base '
               'sin IVA. Sustitúyelo por tus presupuestos reales cuando los '
               'tengas: es COPIA declarada, no un vínculo entre ficheros.')


def hoja_inversion(wb):
    ws = wb.create_sheet(H_INV)
    C.anchos(ws, {'A': 58, 'B': 15, 'C': 12, 'D': 12, 'E': 14, 'F': 80})
    C.encabezar(ws, 'Inversión inicial: cuánto cuesta abrir',
                'Cada línea lleva su columna «¿Lleva IVA?» y su TIPO. Sin esa '
                'columna el CAPEX sale hasta un 21 % desviado, que es el error '
                'más caro de todo el proceso. Lo que no tengas presupuestado, '
                'déjalo en su valor sembrado y márcalo para pedirlo.',
                col_fin='F')
    C.cabecera(ws, I_CAB, [('A', 'Concepto'), ('B', 'Importe (€, base sin IVA)'),
                           ('C', '% s/inversión'), ('D', '¿Lleva IVA?'),
                           ('E', 'Tipo de IVA de la línea'), ('F', 'Notas')])

    refs, fila_listas = C.bloque_listas(
        ws, I_NOTA + 2, [('¿Lleva IVA?', ['Sí', 'No'])], col='A')

    for i, (bloque, partida, importe, base, tipo, fuente, nota_txt) in \
            enumerate(D.CAPEX):
        fila = _FILA_CAPEX[i]
        etiqueta = bloque + ' · ' + partida
        motor.val(ws, 'A%d' % fila, etiqueta, wrap=True)
        if bloque == 'Fondo de maniobra':
            motor.f(ws, 'B%d' % fila,
                    ie('%s*%sC%d/12' % (sup('colchon'), Q_PYG, P['tot_fijos'])),
                    fmt=C.EUR)
        elif importe is None:
            C.entrada(ws, 'B%d' % fila, DEFECTO_CAPEX[i], fmt=C.EUR,
                      etiqueta=etiqueta)
        else:
            C.entrada(ws, 'B%d' % fila, importe, fmt=C.EUR, etiqueta=etiqueta)
        motor.f(ws, 'C%d' % fila, ie('B%d/$B$%d' % (fila, I_TOTAL)), fmt=C.PCT)
        C.entrada(ws, 'D%d' % fila, 'Sí' if tipo else 'No', etiqueta=etiqueta,
                  align='center')
        C.entrada(ws, 'E%d' % fila, tipo, fmt=C.PCT, etiqueta=etiqueta)
        texto = nota_txt or ''
        if importe is None and bloque != 'Fondo de maniobra':
            texto = NOTA_EQUIPO
        if bloque == 'Fondo de maniobra':
            texto = ('Se calcula: meses de colchón x costes fijos mensuales '
                     'del año de crucero. No es inversión que se compra, es '
                     'caja: por eso sale del CAPEX en la fila de abajo.')
        C.nota(ws, 'F%d' % fila, texto + ('' if not fuente else
                                          '  [Fuente: ' + fuente + ']'))
        ws.row_dimensions[fila].height = 30
        if fuente and fuente.startswith('PA-04'):
            C.nota_celda(ws, 'B%d' % fila, 'PA-04')
    C.dv_rango(ws, ['D%d' % _FILA_CAPEX[i] for i in range(len(D.CAPEX))],
               refs['¿Lleva IVA?'], 'Marca Sí o No',
               'Elige «Sí» o «No» de la lista del pie de la hoja.')

    motor.val(ws, 'A%d' % I_SUBTOT, 'SUBTOTAL DE INVERSIÓN (sin el fondo de '
                                    'maniobra)', bold=True)
    motor.f(ws, 'B%d' % I_SUBTOT,
            ie('IF(COUNT(B%d:B%d)=0,"",SUM(B%d:B%d))'
               % (I_INI, I_FIN - 1, I_INI, I_FIN - 1)), fmt=C.EUR, bold=True)
    motor.f(ws, 'C%d' % I_SUBTOT, ie('B%d/$B$%d' % (I_SUBTOT, I_TOTAL)),
            fmt=C.PCT)
    C.total_fila(ws, I_SUBTOT, 'ABCDEF')

    motor.val(ws, 'A%d' % I_FDM, 'FONDO DE MANIOBRA', bold=True)
    motor.f(ws, 'B%d' % I_FDM, ie('B%d' % I_FONDO), fmt=C.EUR, bold=True)
    motor.f(ws, 'C%d' % I_FDM, ie('B%d/$B$%d' % (I_FDM, I_TOTAL)), fmt=C.PCT)
    C.total_fila(ws, I_FDM, 'ABCDEF')

    motor.val(ws, 'A%d' % I_TOTAL, 'INVERSIÓN TOTAL (lo que se compra más la '
                                   'caja del arranque)', bold=True)
    motor.f(ws, 'B%d' % I_TOTAL, ie('B%d+B%d' % (I_SUBTOT, I_FDM)), fmt=C.EUR,
            bold=True)
    C.total_fila(ws, I_TOTAL, 'ABCDEF')

    motor.val(ws, 'A%d' % I_CAPEX, 'CAPEX (inversión sin el fondo de maniobra)')
    motor.f(ws, 'B%d' % I_CAPEX, ie('B%d-B%d' % (I_TOTAL, I_FDM)), fmt=C.EUR)
    C.nota(ws, 'F%d' % I_CAPEX,
           'Es lo que se COMPRA: obra, equipamiento, licencias, packaging y '
           'lanzamiento. El colchón de caja no es inversión, es liquidez.')

    motor.val(ws, 'A%d' % I_IVA,
              'IVA soportado sobre la inversión (recuperable, pero hay que '
              'ADELANTARLO)')
    motor.f(ws, 'B%d' % I_IVA,
            ie('SUMPRODUCT(B%d:B%d,E%d:E%d)'
               % (I_INI, I_FIN, I_INI, I_FIN)), fmt=C.EUR)
    C.nota(ws, 'F%d' % I_IVA,
           'Cada línea a SU tipo, leído de la columna de al lado. Se recupera '
           'por el modelo 303, pero tarda: la hoja de Tesorería lo arrastra '
           'trimestre a trimestre y esa espera es la partida que más sorprende.')

    motor.val(ws, 'A%d' % I_NECES, 'NECESIDAD TOTAL DE CAJA AL ARRANQUE',
              bold=True)
    motor.f(ws, 'B%d' % I_NECES, ie('B%d+B%d' % (I_TOTAL, I_IVA)), fmt=C.EUR,
            bold=True)
    C.total_fila(ws, I_NECES, 'ABCDEF')
    C.nota(ws, 'F%d' % I_NECES,
           'Es la cifra que tiene que cubrir la hoja de Financiación.')

    motor.val(ws, 'A%d' % I_CHK,
              'Comprobación: líneas con «Sí» y tipo cero, o al revés')
    motor.f(ws, 'B%d' % I_CHK,
            ie('SUMPRODUCT((D%d:D%d="Sí")*(E%d:E%d=0))'
               '+SUMPRODUCT((D%d:D%d="No")*(E%d:E%d>0))'
               % (I_INI, I_FIN, I_INI, I_FIN, I_INI, I_FIN, I_INI, I_FIN)),
            fmt=C.ENT)
    motor.semaforo_isnumber(ws, 'B%d:B%d' % (I_CHK, I_CHK), '$B$%d' % I_CHK,
                            operador='>', umbral='0')
    C.nota(ws, 'F%d' % I_CHK,
           'Tiene que valer cero. Si se pone en rojo, una línea dice «Sí» y '
           'tiene el tipo a cero (o al revés) y el IVA de la inversión está '
           'mal calculado.')

    C.seccion(ws, 'A%d' % I_SEC_AM, 'BASES DE AMORTIZACIÓN (no suman a la '
                                    'inversión)')
    motor.val(ws, 'A%d' % I_BASE, 'Inmovilizado amortizable')
    motor.f(ws, 'B%d' % I_BASE,
            ie('B%d-B%d-B%d' % (I_SUBTOT, I_FIANZA, I_PACK)), fmt=C.EUR)
    C.nota(ws, 'F%d' % I_BASE,
           'Todo el subtotal menos la fianza (es un depósito que se recupera) '
           'y el primer pedido de packaging (son existencias, no inmovilizado).')
    motor.val(ws, 'A%d' % I_AMORT, 'Amortización anual del inmovilizado')
    motor.f(ws, 'B%d' % I_AMORT,
            ie('B%d/MAX(1,%s)' % (I_BASE, sup('anios_amort'))), fmt=C.EUR)
    C.nota(ws, 'F%d' % I_AMORT,
           'Es la línea que entra en los costes fijos del P&L. Sin ella la '
           'rentabilidad publicada es mentira: los hornos se gastan.')

    C.parrafo(ws, I_NOTA,
              'Todas las cifras de la columna «Importe» van en BASE SIN IVA. '
              'El IVA de cada línea se calcula aparte, abajo, para que puedas '
              'ver cuánto dinero tienes que adelantar y cuánto vas a '
              'recuperar. La renta que pagues ANTES de abrir no está aquí: si '
              'la vas a pagar, súmala a la línea de fianza y dilo en su nota, '
              'para que el libro de CAPEX y éste sigan diciendo lo mismo.',
              'A', 'F', alto=42)
    C.pagina(ws, titulos='$%d:$%d' % (I_CAB, I_CAB),
             area='A1:F%d' % fila_listas)
    return ws


# ==========================================================================
# Hoja «PyG 3 Años»
# ==========================================================================
def hoja_pyg(wb):
    ws = wb.create_sheet(H_PYG)
    C.anchos(ws, {'A': 58, 'B': 15, 'C': 15, 'D': 15, 'E': 13, 'F': 13,
                  'G': 78})
    C.encabezar(ws, 'Cuenta de resultados a 3 años',
                'El AÑO 2 es el año de crucero: es el que citan la guía y el '
                'business plan. El año 1 lleva la rampa de arranque y el año 3 '
                'crece sobre el crucero. Sueldo del propietario, amortización '
                'e intereses van DENTRO de los costes fijos.', col_fin='G')
    C.cabecera(ws, P['cab'], [('A', 'Concepto'), ('B', 'Año 1 (arranque)'),
                              ('C', 'Año 2 (crucero)'), ('D', 'Año 3'),
                              ('E', '% s/ventas (año 2)'),
                              ('F', 'Tipo de IVA de la línea'), ('G', 'Notas')])

    def sec(fila, texto):
        C.seccion(ws, 'A%d' % fila, texto)
        for letra in 'BCDEFG':
            motor.val(ws, letra + str(fila), '')
            ws[letra + str(fila)].fill = PatternFill('solid',
                                                     fgColor=C.CABECERA)

    def pct(fila):
        motor.f(ws, 'E%d' % fila,
                ie('C%d/$C$%d' % (fila, P['ingresos'])), fmt=C.PCT)

    G = 'IF(B%d="","",%s)'

    sec(P['sec_ing'], 'INGRESOS')
    motor.val(ws, 'A%d' % P['tickets'], 'Tickets al día')
    motor.f(ws, 'B%d' % P['tickets'], ie(sup('tickets')), fmt=C.ENT)
    motor.f(ws, 'C%d' % P['tickets'], ie(sup('tickets')), fmt=C.ENT)
    motor.f(ws, 'D%d' % P['tickets'],
            ie('C%d*(1+%s)' % (P['tickets'], sup('crec_a3'))), fmt=C.ENT)
    C.nota(ws, 'G%d' % P['tickets'],
           'Los años 1 y 2 comparten el tráfico de crucero; lo que separa al '
           'año 1 es la rampa de la fila de abajo, no menos clientes al día en '
           'el mes 12.')

    motor.val(ws, 'A%d' % P['ticket'], 'Ticket medio sin IVA (€)')
    for col in 'BCD':
        motor.f(ws, '%s%d' % (col, P['ticket']), ie(sup('ticket_sin')),
                fmt=C.EUR)
    C.nota(ws, 'G%d' % P['ticket'],
           'Constante en euros del año 1: una subida de precios es una '
           'decisión aparte, y se toma en la carta, no aquí.')

    motor.val(ws, 'A%d' % P['dias'], 'Días de apertura al año')
    for col in 'BCD':
        motor.f(ws, '%s%d' % (col, P['dias']), ie(sup('dias')), fmt=C.ENT)

    motor.val(ws, 'A%d' % P['actividad'],
              'Actividad sobre la velocidad de crucero (%)')
    motor.f(ws, 'B%d' % P['actividad'], ie(sup('factor_a1')), fmt=C.PCT)
    C.entrada(ws, 'C%d' % P['actividad'], 1.0, fmt=C.PCT,
              etiqueta='Actividad del año 2 sobre el crucero')
    C.entrada(ws, 'D%d' % P['actividad'], 1.0, fmt=C.PCT,
              etiqueta='Actividad del año 3 sobre el crucero')
    C.nota(ws, 'G%d' % P['actividad'],
           'El año 1 lo calcula «0. Supuestos» desde la rampa y la '
           'estacionalidad de la hoja de Tesorería. Los años 2 y 3 están al '
           '100 % porque el año 2 es, por definición, el de crucero: si crees '
           'que vas a tardar más en llegar, bájalos.')

    motor.val(ws, 'A%d' % P['ingresos'], 'INGRESOS TOTALES (sin IVA)',
              bold=True)
    for col in 'BCD':
        motor.f(ws, '%s%d' % (col, P['ingresos']),
                ie('IF({0}{1}*{0}{2}*{0}{3}*{0}{4}=0,"",{0}{1}*{0}{2}*{0}{3}'
                   '*{0}{4})'.format(col, P['tickets'], P['ticket'],
                                     P['dias'], P['actividad'])),
                fmt=C.EUR, bold=True)
    pct(P['ingresos'])
    C.total_fila(ws, P['ingresos'], 'ABCDEFG')

    motor.val(ws, 'A%d' % P['v_past'],
              'Ventas de pastelería, bollería y tartas')
    motor.val(ws, 'A%d' % P['v_pan'], 'Ventas de pan de acompañamiento')
    for col in 'BCD':
        motor.f(ws, '%s%d' % (col, P['v_past']),
                ie('IF({0}{1}="","",{0}{1}*(1-{2}))'
                   .format(col, P['ingresos'], sup('peso_pan'))), fmt=C.EUR)
        motor.f(ws, '%s%d' % (col, P['v_pan']),
                ie('IF({0}{1}="","",{0}{1}*{2})'
                   .format(col, P['ingresos'], sup('peso_pan'))), fmt=C.EUR)
    pct(P['v_past'])
    pct(P['v_pan'])
    motor.f(ws, 'F%d' % P['v_past'], ie(sup('iva_prod')), fmt=C.PCT)
    motor.f(ws, 'F%d' % P['v_pan'], ie(sup('iva_pan')), fmt=C.PCT)
    C.nota_celda(ws, 'F%d' % P['v_past'], 'PA-36')
    C.nota_celda(ws, 'F%d' % P['v_pan'], 'PA-36b')
    C.nota(ws, 'G%d' % P['v_pan'],
           'El pan va al 4 % y todo lo demás al 10 %. Separar las dos líneas '
           'no es un capricho contable: es lo que hace que la caja y el '
           'modelo 303 cuadren.')

    sec(P['sec_var'], 'COSTES VARIABLES')
    motor.val(ws, 'A%d' % P['coste_ventas'],
              'Coste de ventas (materia prima, packaging y merma)')
    for col in 'BCD':
        motor.f(ws, '%s%d' % (col, P['coste_ventas']),
                ie('IF({0}{1}="","",{0}{1}*{2})'
                   .format(col, P['ingresos'], sup('food_cost'))), fmt=C.EUR)
    pct(P['coste_ventas'])
    motor.f(ws, 'F%d' % P['coste_ventas'], ie(sup('iva_compras')), fmt=C.PCT)
    C.nota(ws, 'G%d' % P['coste_ventas'],
           'A la REGLA ÚNICA del 32 %, no al food cost de escandallo de la '
           'carta: la regla incluye el producto de terceros que revende el '
           'despacho, la merma real, el packaging y la subida de la materia '
           'prima. Es el supuesto conservador, y la diferencia con tu '
           'escandallo es tu colchón.')

    motor.val(ws, 'A%d' % P['tot_var'], 'TOTAL COSTES VARIABLES', bold=True)
    for col in 'BCD':
        motor.f(ws, '%s%d' % (col, P['tot_var']),
                ie('IF({0}{1}="","",SUM({0}{2}:{0}{2}))'
                   .format(col, P['ingresos'], P['coste_ventas'])),
                fmt=C.EUR, bold=True)
    pct(P['tot_var'])
    C.total_fila(ws, P['tot_var'], 'ABCDEFG')

    motor.val(ws, 'A%d' % P['margen'], 'MARGEN BRUTO', bold=True)
    for col in 'BCD':
        motor.f(ws, '%s%d' % (col, P['margen']),
                ie('IF({0}{1}="","",{0}{1}-{0}{2})'
                   .format(col, P['ingresos'], P['tot_var'])),
                fmt=C.EUR, bold=True)
    pct(P['margen'])
    C.total_fila(ws, P['margen'], 'ABCDEFG')

    sec(P['sec_fijos'], 'COSTES FIJOS')
    motor.val(ws, 'A%d' % P['personal'],
              'Personal (nóminas y Seguridad Social)')
    for col in 'BCD':
        motor.f(ws, '%s%d' % (col, P['personal']),
                ie('IF({0}{1}="","",{2}$I${3})'
                   .format(col, P['ingresos'], Q_PER, R_TOT)), fmt=C.EUR)
    motor.val(ws, 'F%d' % P['personal'], 0.0, fmt=C.PCT)
    C.nota(ws, 'G%d' % P['personal'],
           'Sale de la hoja de Personal: es el MISMO número, no una estimación '
           'aparte. Las nóminas no llevan IVA. Si contratas a alguien en el '
           'año 2 o el 3, añádelo allí y aquí se ve.')

    for clave, etiqueta, sup_clave, meses in (
            ('propietario',
             'Retribución del propietario (incluida la cuota de autónomos)',
             'retribucion', 12),
            ('alquiler', 'Alquiler del local', 'renta', 12),
            ('suministros', 'Suministros (electricidad, gas y agua)',
             'suministros', 12)):
        fila = P[clave]
        motor.val(ws, 'A%d' % fila, etiqueta)
        motor.f(ws, 'B%d' % fila,
                ie('IF(B%d="","",%s*%d)' % (P['ingresos'], sup(sup_clave),
                                            meses)), fmt=C.EUR)
        motor.f(ws, 'C%d' % fila,
                ie('IF(C%d="","",B%d*(1+%s))'
                   % (P['ingresos'], fila, sup('subida'))), fmt=C.EUR)
        motor.f(ws, 'D%d' % fila,
                ie('IF(D%d="","",C%d*(1+%s))'
                   % (P['ingresos'], fila, sup('subida'))), fmt=C.EUR)
        pct(fila)
    motor.val(ws, 'F%d' % P['propietario'], 0.0, fmt=C.PCT)
    motor.f(ws, 'F%d' % P['alquiler'], ie(sup('iva_gen')), fmt=C.PCT)
    motor.f(ws, 'F%d' % P['suministros'], ie(sup('iva_gen')), fmt=C.PCT)
    C.nota(ws, 'G%d' % P['propietario'],
           'Renglón propio, y no está dentro de la plantilla de cinco '
           'personas. Un negocio que sólo es rentable porque su dueño no cobra '
           'no es rentable.')
    C.nota(ws, 'G%d' % P['alquiler'],
           'El importe mensual está en «0. Supuestos»; aquí se multiplica por '
           'doce. El arrendamiento de local de negocio lleva IVA al tipo '
           'general.')
    C.nota(ws, 'G%d' % P['suministros'],
           'Luz y gas van al 21 %; el agua, al 10 %. El libro aplica a toda la '
           'partida el tipo de la columna de al lado: si el agua te pesa, '
           'sepárala en su propia línea.')

    for clave, etiqueta, valor, tipo, nota_txt in FIJOS_VERDES:
        fila = P[clave]
        motor.val(ws, 'A%d' % fila, etiqueta)
        C.entrada(ws, 'B%d' % fila, valor, fmt=C.EUR, etiqueta=etiqueta)
        motor.f(ws, 'C%d' % fila,
                ie('IF(C%d="","",B%d*(1+%s))'
                   % (P['ingresos'], fila, sup('subida'))), fmt=C.EUR)
        motor.f(ws, 'D%d' % fila,
                ie('IF(D%d="","",C%d*(1+%s))'
                   % (P['ingresos'], fila, sup('subida'))), fmt=C.EUR)
        pct(fila)
        C.entrada(ws, 'F%d' % fila, tipo, fmt=C.PCT, etiqueta=etiqueta)
        C.nota(ws, 'G%d' % fila, nota_txt)

    motor.val(ws, 'A%d' % P['amortizacion'], 'Amortización del inmovilizado')
    for col in 'BCD':
        motor.f(ws, '%s%d' % (col, P['amortizacion']),
                ie('IF({0}{1}="","",{2}$B${3})'
                   .format(col, P['ingresos'], Q_INV, I_AMORT)), fmt=C.EUR)
    pct(P['amortizacion'])
    motor.val(ws, 'F%d' % P['amortizacion'], 0.0, fmt=C.PCT)
    C.nota(ws, 'G%d' % P['amortizacion'],
           'Sale de la base de amortización de la hoja de Inversión. No sube '
           'con el IPC: es un reparto contable de una inversión ya hecha.')

    motor.val(ws, 'A%d' % P['financieros'], 'Gastos financieros del préstamo')
    for col, anio in (('B', 0), ('C', 1), ('D', 2)):
        motor.f(ws, '%s%d' % (col, P['financieros']),
                ie('IF({0}{1}="","",{2}$B${3})'
                   .format(col, P['ingresos'], Q_FIN, F_ANIO_INI + anio)),
                fmt=C.EUR)
    pct(P['financieros'])
    motor.val(ws, 'F%d' % P['financieros'], 0.0, fmt=C.PCT)
    C.nota(ws, 'G%d' % P['financieros'],
           'Sólo los INTERESES son gasto. La devolución del principal sale de '
           'la caja pero no del resultado: va en la hoja de Tesorería.')

    motor.val(ws, 'A%d' % P['tot_fijos'], 'TOTAL COSTES FIJOS', bold=True)
    for col in 'BCD':
        motor.f(ws, '%s%d' % (col, P['tot_fijos']),
                ie('IF({0}{1}="","",SUM({0}{2}:{0}{3}))'
                   .format(col, P['ingresos'], P['personal'],
                           P['financieros'])), fmt=C.EUR, bold=True)
    pct(P['tot_fijos'])
    C.total_fila(ws, P['tot_fijos'], 'ABCDEFG')

    motor.val(ws, 'A%d' % P['iva_var'], 'IVA soportado de los costes variables')
    for col in 'BCD':
        motor.f(ws, '%s%d' % (col, P['iva_var']),
                ie('IF({0}{1}="","",{0}{2}*$F${2})'
                   .format(col, P['ingresos'], P['coste_ventas'])), fmt=C.EUR)
    motor.val(ws, 'A%d' % P['iva_fijos'], 'IVA soportado de los costes fijos')
    for col in 'BCD':
        motor.f(ws, '%s%d' % (col, P['iva_fijos']),
                ie('IF({0}{1}="","",SUMPRODUCT({0}{2}:{0}{3},$F${2}:$F${3}))'
                   .format(col, P['ingresos'], P['personal'],
                           P['financieros'])), fmt=C.EUR)
    C.nota(ws, 'G%d' % P['iva_fijos'],
           'Cada línea a SU tipo, leído de la columna de al lado. Es lo que la '
           'hoja de Tesorería paga de más a los proveedores y luego recupera '
           'en el modelo 303.')

    motor.val(ws, 'A%d' % P['rai'], 'RESULTADO ANTES DE IMPUESTOS', bold=True)
    for col in 'BCD':
        motor.f(ws, '%s%d' % (col, P['rai']),
                ie('IF({0}{1}="","",{0}{2}-{0}{3})'
                   .format(col, P['ingresos'], P['margen'], P['tot_fijos'])),
                fmt=C.EUR, bold=True)
    pct(P['rai'])
    C.total_fila(ws, P['rai'], 'ABCDEFG')

    motor.val(ws, 'A%d' % P['bases_ini'], 'Bases negativas pendientes al inicio')
    motor.f(ws, 'B%d' % P['bases_ini'],
            ie('IF(B%d="","",%s)' % (P['ingresos'], sup('bases_neg'))),
            fmt=C.EUR)
    motor.f(ws, 'C%d' % P['bases_ini'],
            ie('IF(C%d="","",B%d)' % (P['ingresos'], P['bases_fin'])),
            fmt=C.EUR)
    motor.f(ws, 'D%d' % P['bases_ini'],
            ie('IF(D%d="","",C%d)' % (P['ingresos'], P['bases_fin'])),
            fmt=C.EUR)
    C.nota(ws, 'G%d' % P['bases_ini'],
           'Art. 26 de la Ley del Impuesto sobre Sociedades: las pérdidas de un '
           'ejercicio se compensan con los beneficios de los siguientes.')

    motor.val(ws, 'A%d' % P['base_imp'], 'Base imponible (después de compensar)')
    for col in 'BCD':
        motor.f(ws, '%s%d' % (col, P['base_imp']),
                ie('IF({0}{1}="","",MAX(0,{0}{2}-{0}{3}))'
                   .format(col, P['ingresos'], P['rai'], P['bases_ini'])),
                fmt=C.EUR)
    motor.val(ws, 'A%d' % P['ejercicios'],
              'Ejercicios con base positiva (acumulado)')
    motor.f(ws, 'B%d' % P['ejercicios'],
            ie('IF(B%d="","",IF(B%d>0,1,0))'
               % (P['ingresos'], P['base_imp'])), fmt=C.ENT)
    for col, ant in (('C', 'B'), ('D', 'C')):
        motor.f(ws, '%s%d' % (col, P['ejercicios']),
                ie('IF({0}{1}="","",{2}{3}+IF({0}{4}>0,1,0))'
                   .format(col, P['ingresos'], ant, P['ejercicios'],
                           P['base_imp'])), fmt=C.ENT)
    motor.val(ws, 'A%d' % P['tipo_is'],
              'Tipo de Impuesto de Sociedades aplicado')
    for col in 'BCD':
        motor.f(ws, '%s%d' % (col, P['tipo_is']),
                ie('IF({0}{1}="","",IF({0}{2}<=0,"",IF({0}{3}<=2,{4},{5})))'
                   .format(col, P['ingresos'], P['base_imp'], P['ejercicios'],
                           sup('is_nueva'), sup('is_gen'))), fmt=C.PCT)
    C.nota(ws, 'G%d' % P['tipo_is'],
           'El 15 % de entidad de nueva creación se aplica al PRIMER ejercicio '
           'con base positiva y al siguiente, no a los dos primeros años '
           'naturales. Si tributas en IRPF, esta línea no te aplica.')
    motor.val(ws, 'A%d' % P['is'], 'Impuesto de Sociedades')
    for col in 'BCD':
        motor.f(ws, '%s%d' % (col, P['is']),
                ie('IF({0}{1}="","",IF({0}{2}="",0,{0}{3}*{0}{2}))'
                   .format(col, P['ingresos'], P['tipo_is'], P['base_imp'])),
                fmt=C.EUR)
    motor.val(ws, 'A%d' % P['bases_fin'],
              'Bases negativas pendientes al cierre')
    for col in 'BCD':
        motor.f(ws, '%s%d' % (col, P['bases_fin']),
                ie('IF({0}{1}="","",MAX(0,{0}{2}-{0}{3}))'
                   .format(col, P['ingresos'], P['bases_ini'], P['rai'])),
                fmt=C.EUR)
    motor.val(ws, 'A%d' % P['neto'], 'RESULTADO NETO', bold=True)
    for col in 'BCD':
        motor.f(ws, '%s%d' % (col, P['neto']),
                ie('IF({0}{1}="","",{0}{2}-{0}{3})'
                   .format(col, P['ingresos'], P['rai'], P['is'])),
                fmt=C.EUR, bold=True)
    pct(P['neto'])
    C.total_fila(ws, P['neto'], 'ABCDEFG')
    motor.semaforo_isnumber(ws, 'B%d:D%d' % (P['neto'], P['neto']),
                            'B%d' % P['neto'], operador='<', umbral='0')

    # --- ratios -----------------------------------------------------------
    sec(P['sec_ratios'], 'RATIOS CLAVE')
    C.cabecera(ws, P['cab_ratios'],
               [('A', 'Ratio'), ('B', 'Año 1'), ('C', 'Año 2'), ('D', 'Año 3'),
                ('E', 'Umbral'), ('F', ''), ('G', 'Comentario')], altura=22)
    RATIOS = (
        ('r_margen', 'Margen bruto / ventas', P['margen'], 'margen_obj',
         'El suelo es la regla única de margen del propio libro. Por debajo, o '
         'suben los precios o baja el coste de la materia prima.'),
        ('r_coste', 'Coste de ventas / ventas', P['coste_ventas'], 'food_cost',
         'Es el food cost proyectado. Cuando midas el tuyo con el escandallo, '
         'cámbialo en «0. Supuestos» y mira qué le pasa a este cuadro.'),
        ('r_personal', 'Coste de personal / ventas', P['personal'],
         'techo_personal',
         'En un obrador con despacho el personal es la primera partida. Por '
         'encima del techo, o subes precios o redimensionas turnos.'),
        ('r_alquiler', 'Alquiler / ventas', P['alquiler'], 'techo_alquiler',
         'Un alquiler por encima del techo se come el margen de todo lo demás, '
         'y no se renegocia una vez firmado.'),
        ('r_neto', 'Resultado neto / ventas', P['neto'], 'neto_suelo',
         'El suelo es el extremo bajo de la banda de referencia declarada por '
         'un pastelero de cuarta generación sobre su propia casa (8-12 %).'),
    )
    for clave, etiqueta, fila_origen, umbral, nota_txt in RATIOS:
        fila = P[clave]
        motor.val(ws, 'A%d' % fila, etiqueta)
        for col in 'BCD':
            motor.f(ws, '%s%d' % (col, fila),
                    ie('IF({0}{1}="","",{0}{2}/{0}{1})'
                       .format(col, P['ingresos'], fila_origen)), fmt=C.PCT)
        motor.f(ws, 'E%d' % fila, ie(sup(umbral)), fmt=C.PCT)
        C.nota(ws, 'G%d' % fila, nota_txt)

    motor.val(ws, 'A%d' % P['r_banda'],
              'Margen neto dentro de la banda de referencia (8-12 %)')
    for col in 'BCD':
        motor.f(ws, '%s%d' % (col, P['r_banda']),
                ie('IF({0}{1}="","",IF(AND({0}{1}>={2},{0}{1}<={3}),"Sí",'
                   '"No: revísalo"))'.format(col, P['r_neto'],
                                             sup('neto_suelo'),
                                             sup('neto_techo'))))
    C.nota(ws, 'G%d' % P['r_banda'],
           'Si sale por ENCIMA del 12 %, desconfía antes de enseñarlo: casi '
           'siempre falta un coste. Si sale por debajo del 8 %, el negocio no '
           'está mal planteado por fuerza, pero tiene poco donde caerse.')

    motor.val(ws, 'A%d' % P['r_equilibrio'], 'Punto de equilibrio alcanzado')
    for col, col_pe in (('B', 'B'), ('C', 'C'), ('D', 'D')):
        motor.f(ws, '%s%d' % (col, P['r_equilibrio']),
                ie('IF({0}{1}="","",IF({0}{1}>=\'{2}\'!${3}${4},"Sí","No"))'
                   .format(col, P['ingresos'], H_PEQ, col_pe, E['ing_anio'])))
    C.nota(ws, 'G%d' % P['r_equilibrio'],
           'Compara las ventas de CADA año con el umbral contable de ese mismo '
           'año, que calcula la hoja «Punto de Equilibrio». El umbral de caja, '
           'que es el que de verdad te deja dormir, está allí.')

    C.parrafo(ws, P['nota'],
              'Todas las cifras van SIN IVA. Las columnas de los años 2 y 3 '
              'están en euros del año 1 mientras la «Subida anual de los '
              'costes fijos» de «0. Supuestos» valga cero. Y una advertencia '
              'que vale dinero: si tu proyección del año 1 sale en positivo, '
              'míralo dos veces. Casi ninguna apertura gana dinero el primer '
              'año, y lo que decide si sobrevives no es este cuadro sino la '
              'hoja de Tesorería.', 'A', 'G', alto=44)
    C.pagina(ws, titulos='$%d:$%d' % (P['cab'], P['cab']),
             area='A1:G%d' % P['nota'])
    return ws


# ==========================================================================
# Hoja «Punto de Equilibrio»
# ==========================================================================
def hoja_equilibrio(wb):
    ws = wb.create_sheet(H_PEQ)
    C.anchos(ws, {'A': 56, 'B': 15, 'C': 15, 'D': 15, 'E': 80})
    C.encabezar(ws, 'Punto de equilibrio: cuántos tickets al día',
                'Dos umbrales, y no dicen lo mismo. El CONTABLE incluye la '
                'amortización; el de CAJA le quita la amortización (que no se '
                'paga) y le suma la devolución de principal (que sí). El '
                'segundo es el que decide si llegas a fin de mes.', col_fin='E')
    C.cabecera(ws, E['cab'], [('A', 'Variable'), ('B', 'Año 1'),
                              ('C', 'Año 2'), ('D', 'Año 3'), ('E', 'Notas')])

    def sec(fila, texto):
        C.seccion(ws, 'A%d' % fila, texto)
        for letra in 'BCDE':
            motor.val(ws, letra + str(fila), '')
            ws[letra + str(fila)].fill = PatternFill('solid',
                                                     fgColor=C.CABECERA)

    def linea(fila, etiqueta, plantilla, fmt, nota_txt=None, bold=None):
        motor.val(ws, 'A%d' % fila, etiqueta, bold=bold)
        for col in 'BCD':
            motor.f(ws, '%s%d' % (col, fila), plantilla(col), fmt=fmt,
                    bold=bold)
        if nota_txt:
            C.nota(ws, 'E%d' % fila, nota_txt)

    guardia = 'IF({0}{1}{2}="","",{3})'

    def g(col, expr):
        return ie(guardia.format(Q_PYG, col, P['ingresos'], expr))

    sec(E['sec_base'], 'DATOS BASE POR AÑO')
    linea(E['fijos_anio'], 'Costes fijos anuales',
          lambda c: g(c, '%s%s%d' % (Q_PYG, c, P['tot_fijos'])), C.EUR,
          'Los MISMOS costes fijos del P&L. Cada año tiene los suyos: por eso '
          'el umbral no puede ser uno solo para los tres.')
    linea(E['fijos_mes'], 'Costes fijos mensuales',
          lambda c: g(c, '%s%d/12' % (c, E['fijos_anio'])), C.EUR)
    linea(E['ticket'], 'Ticket medio sin IVA (€)',
          lambda c: g(c, '%s%s%d' % (Q_PYG, c, P['ticket'])), C.EUR)
    linea(E['cv_ticket'], 'Coste variable por ticket (€)',
          lambda c: g(c, '{0}{1}{2}/({0}{1}{3}*{0}{1}{4}*{0}{1}{5})'
                      .format(Q_PYG, c, P['tot_var'], P['tickets'],
                              P['dias'], P['actividad'])), C.EUR,
          'Sale del total de costes variables del P&L dividido por los tickets '
          'de ese año, no de una estimación aparte.')
    linea(E['mc_ticket'], 'Margen de contribución por ticket (€)',
          lambda c: g(c, '%s%d-%s%d' % (c, E['ticket'], c, E['cv_ticket'])),
          C.EUR, 'Ticket menos coste variable unitario. Es lo que deja cada '
          'cliente para pagar los fijos.')
    linea(E['dias'], 'Días de apertura al año',
          lambda c: g(c, '%s%s%d' % (Q_PYG, c, P['dias'])), C.ENT)
    for col, anio in (('B', 0), ('C', 1), ('D', 2)):
        motor.f(ws, '%s%d' % (col, E['principal']),
                g(col, '%s$C$%d' % (Q_FIN, F_ANIO_INI + anio)), fmt=C.EUR)
    motor.val(ws, 'A%d' % E['principal'],
              'Devolución de principal del año (sale de caja)')
    C.nota(ws, 'E%d' % E['principal'],
           'El principal NO es gasto del P&L, pero sale de la caja. Sin él, el '
           'umbral de equilibrio se publica por debajo de lo que de verdad hay '
           'que facturar.')

    sec(E['sec_cont'], 'PUNTO DE EQUILIBRIO CONTABLE (incluye la amortización)')
    linea(E['tk_anio'], 'Tickets necesarios al año',
          lambda c: g(c, 'IF({0}{1}<=0,"",{0}{2}/{0}{1})'
                      .format(c, E['mc_ticket'], E['fijos_anio'])), C.ENT)
    linea(E['tk_dia'], 'Tickets necesarios al día',
          lambda c: g(c, '%s%d/%s%d' % (c, E['tk_anio'], c, E['dias'])),
          C.DEC1)
    linea(E['ing_anio'], 'Ingresos necesarios al año (sin IVA)',
          lambda c: g(c, '%s%d*%s%d' % (c, E['tk_anio'], c, E['ticket'])),
          C.EUR, bold=True)
    linea(E['ing_mes'], 'Ingresos necesarios al mes (sin IVA)',
          lambda c: g(c, '%s%d/12' % (c, E['ing_anio'])), C.EUR, bold=True)
    C.total_fila(ws, E['ing_anio'], 'ABCDE')
    C.total_fila(ws, E['ing_mes'], 'ABCDE')

    sec(E['sec_caja'], 'PUNTO DE EQUILIBRIO DE CAJA (incluye la cuota del '
                       'préstamo)')
    linea(E['fijos_caja'], 'Costes fijos de caja más el principal del año',
          lambda c: g(c, '{0}{1}-{2}{0}{3}+{0}{4}'
                      .format(c, E['fijos_anio'], Q_PYG, P['amortizacion'],
                              E['principal'])), C.EUR,
          'Los fijos del P&L sin la amortización (que no se paga) y CON la '
          'devolución de principal (que sí). Es el umbral que decide si '
          'sobrevives, no el contable.')
    linea(E['tkc_anio'], 'Tickets necesarios al año (caja)',
          lambda c: g(c, 'IF({0}{1}<=0,"",{0}{2}/{0}{1})'
                      .format(c, E['mc_ticket'], E['fijos_caja'])), C.ENT)
    linea(E['tkc_dia'], 'Tickets necesarios al día (caja)',
          lambda c: g(c, '%s%d/%s%d' % (c, E['tkc_anio'], c, E['dias'])),
          C.DEC1, bold=True)
    linea(E['ingc_anio'], 'Ingresos necesarios al año (caja, sin IVA)',
          lambda c: g(c, '%s%d*%s%d' % (c, E['tkc_anio'], c, E['ticket'])),
          C.EUR)
    C.total_fila(ws, E['tkc_dia'], 'ABCDE')

    sec(E['sec_contraste'], 'CONTRASTE CON EL PLAN')
    linea(E['tk_previstos'], 'Tickets al día previstos en el plan',
          lambda c: g(c, '{0}{1}{2}*{0}{1}{3}'
                      .format(Q_PYG, c, P['tickets'], P['actividad'])),
          C.DEC1)
    linea(E['holgura'], 'Holgura sobre el equilibrio de CAJA (%)',
          lambda c: g(c, 'IF({0}{1}<=0,"",{0}{2}/{0}{1}-1)'
                      .format(c, E['tkc_dia'], E['tk_previstos'])), C.PCT,
          'Cuánto puedes caer sobre lo previsto antes de no poder pagar los '
          'fijos y la cuota. El mínimo exigible está en «0. Supuestos».')
    linea(E['semaforo'], 'Lectura de la holgura',
          lambda c: g(c, 'IF({0}{1}<0,"Por debajo del equilibrio",'
                      'IF({0}{1}<{2},"Ajustada","Suficiente"))'
                      .format(c, E['holgura'], sup('holgura'))), None)
    motor.semaforo_isnumber(ws, 'B%d:D%d' % (E['holgura'], E['holgura']),
                            'B%d' % E['holgura'], operador='<', umbral='0')
    C.parrafo(ws, E['nota'],
              'Lectura rápida: si los tickets al día previstos están por '
              'debajo de los tickets al día de CAJA, el negocio no genera para '
              'pagar los fijos y la cuota, y da igual lo que diga el '
              'resultado contable. Si están justo por encima, no tienes margen '
              'para un mes malo: negocia carencia, baja inversión o revisa el '
              'ticket medio antes de firmar nada.', 'A', 'E', alto=40)

    sec(E['sec_sens'], 'SENSIBILIDAD DEL PUNTO DE EQUILIBRIO (año de crucero)')
    motor.val(ws, 'A%d' % E['sens_cab'],
              'Tickets al día necesarios: ticket medio (columnas) x coste '
              'variable por ticket (filas)', wrap=True)
    ws.row_dimensions[E['sens_cab']].height = 30
    for i, col in enumerate(('B', 'C', 'D', 'E')):
        desp = i - 1
        if desp == 0:
            expr = '$C$%d' % E['ticket']
        elif desp < 0:
            expr = '$C$%d-1' % E['ticket']
        else:
            expr = '$C$%d+%d' % (E['ticket'], desp)
        motor.f(ws, '%s%d' % (col, E['sens_cab']), ie(expr), fmt=C.EUR,
                bold=True)
    for j in range(5):
        fila = E['sens_ini'] + j
        desp = j - 2
        if desp == 0:
            expr = '$C$%d' % E['cv_ticket']
        elif desp < 0:
            expr = '$C$%d-%d*%s' % (E['cv_ticket'], -desp, sup('paso_sens'))
        else:
            expr = '$C$%d+%d*%s' % (E['cv_ticket'], desp, sup('paso_sens'))
        motor.f(ws, 'A%d' % fila, ie(expr), fmt=C.EUR, bold=True)
        for col in ('B', 'C', 'D', 'E'):
            motor.f(ws, '%s%d' % (col, fila),
                    ie('IF({0}${1}-$A{2}<=0,"",$C${3}/({0}${1}-$A{2})/$C${4})'
                       .format(col, E['sens_cab'], fila, E['fijos_anio'],
                               E['dias'])), fmt=C.DEC1)
    C.parrafo(ws, E['sens_nota'],
              'La celda del centro es tu escenario. Muévete una columna a la '
              'izquierda y estarás viendo qué pasa si tienes que bajar el '
              'ticket medio un euro; una fila hacia abajo, qué pasa si la '
              'mantequilla y el cacao te suben el coste variable. Los dos '
              'movimientos han pasado en los últimos tres años.', 'A', 'E',
              alto=40)
    C.pagina(ws, titulos='$%d:$%d' % (E['cab'], E['cab']),
             area='A1:E%d' % E['sens_nota'])
    return ws


# ==========================================================================
# Hoja «Escenarios»
# ==========================================================================
def hoja_escenarios(wb):
    ws = wb.create_sheet(H_ESC)
    C.anchos(ws, {'A': 54, 'B': 16, 'C': 16, 'D': 16, 'E': 80})
    C.encabezar(ws, 'Escenarios: pesimista, realista y optimista',
                'La columna «Realista» LEE el año de crucero del P&L, así que '
                'no puede desviarse de él. Las otras dos son verdes: son tus '
                'apuestas, y la de la izquierda es la que hay que poder '
                'sobrevivir.', col_fin='E')
    C.cabecera(ws, X['cab'], [('A', 'Métrica'), ('B', 'Pesimista'),
                              ('C', 'Realista'), ('D', 'Optimista'),
                              ('E', 'Notas')])

    ENTRADAS = ((X['tickets'], 'Tickets al día', 140, 220, C.ENT,
                 P['tickets'],
                 'SUPUESTOS los dos extremos. El pesimista no es «un mal día»: '
                 'es el año entero yendo mal.'),
                (X['ticket'], 'Ticket medio sin IVA (€)', 5.20, 6.60, C.EUR,
                 P['ticket'],
                 'SUPUESTOS. Bajar el ticket medio un euro pesa más que perder '
                 'veinte clientes al día: míralo en la tabla de sensibilidad '
                 'del punto de equilibrio.'),
                (X['dias'], 'Días de apertura al año', 290, 305, C.ENT,
                 P['dias'], 'SUPUESTOS.'))
    for fila, etiqueta, pes, opt, fmt, fila_pyg, nota_txt in ENTRADAS:
        motor.val(ws, 'A%d' % fila, etiqueta)
        C.entrada(ws, 'B%d' % fila, pes, fmt=fmt, etiqueta=etiqueta + ' (pes.)')
        motor.f(ws, 'C%d' % fila, ie('%sC%d' % (Q_PYG, fila_pyg)), fmt=fmt)
        C.entrada(ws, 'D%d' % fila, opt, fmt=fmt, etiqueta=etiqueta + ' (opt.)')
        C.nota(ws, 'E%d' % fila, nota_txt)

    def linea(fila, etiqueta, plantilla, fmt, nota_txt=None, bold=None,
              solo=None):
        motor.val(ws, 'A%d' % fila, etiqueta, bold=bold)
        for col in (solo or 'BCD'):
            motor.f(ws, '%s%d' % (col, fila), plantilla(col), fmt=fmt,
                    bold=bold)
        if nota_txt:
            C.nota(ws, 'E%d' % fila, nota_txt)

    def g(col, expr):
        return ie('IF({0}{1}="","",{2})'.format(col, X['ingresos'], expr))

    linea(X['ingresos'], 'INGRESOS ANUALES (sin IVA)',
          lambda c: ie('IF({0}{1}*{0}{2}*{0}{3}=0,"",{0}{1}*{0}{2}*{0}{3})'
                       .format(c, X['tickets'], X['ticket'], X['dias'])),
          C.EUR, bold=True)
    C.total_fila(ws, X['ingresos'], 'ABCDE')
    linea(X['coste'], 'Coste de ventas (regla única del food cost)',
          lambda c: g(c, '%s%d*%s' % (c, X['ingresos'], sup('food_cost'))),
          C.EUR)
    linea(X['tot_var'], 'TOTAL COSTES VARIABLES',
          lambda c: g(c, 'SUM({0}{1}:{0}{1})'.format(c, X['coste'])), C.EUR)
    linea(X['margen'], 'MARGEN BRUTO',
          lambda c: g(c, '%s%d-%s%d' % (c, X['ingresos'], c, X['tot_var'])),
          C.EUR, bold=True)
    linea(X['fijos'], 'COSTES FIJOS (los del año de crucero del P&L)',
          lambda c: g(c, '%s$C$%d' % (Q_PYG, P['tot_fijos'])), C.EUR,
          'Los costes fijos no cambian con el escenario: por eso son fijos. Si '
          'en el pesimista se disparan sobre las ventas, el ajuste hay que '
          'hacerlo en la hoja de Personal, no aquí.')
    linea(X['rai'], 'RESULTADO ANTES DE IMPUESTOS',
          lambda c: g(c, '%s%d-%s%d' % (c, X['margen'], c, X['fijos'])),
          C.EUR, bold=True)
    linea(X['is'], 'Impuesto de Sociedades',
          lambda c: g(c, 'MAX(0,%s%d-%s)*%s'
                      % (c, X['rai'], sup('bases_neg'), sup('is_nueva'))),
          C.EUR,
          'Al tipo de entidad de nueva creación, compensando las bases '
          'negativas anteriores. Con resultado negativo no hay cuota.')
    linea(X['neto'], 'RESULTADO NETO',
          lambda c: g(c, '%s%d-%s%d' % (c, X['rai'], c, X['is'])), C.EUR,
          bold=True)
    C.total_fila(ws, X['neto'], 'ABCDE')
    motor.semaforo_isnumber(ws, 'B%d:D%d' % (X['neto'], X['neto']),
                            'B%d' % X['neto'], operador='<', umbral='0')
    linea(X['margen_neto'], 'Margen neto (%)',
          lambda c: g(c, '%s%d/%s%d' % (c, X['neto'], c, X['ingresos'])),
          C.PCT,
          'Contrástalo con la banda de referencia del 8-12 % de '
          '«0. Supuestos». El optimista que se va muy por encima suele estar '
          'olvidándose de un coste, no descubriendo un negocio mejor.')
    linea(X['tk_equilibrio'], 'Tickets al día para el equilibrio',
          lambda c: g(c, 'IF({0}{1}-{0}{2}/({0}{3}*{0}{4})<=0,"",'
                      '{0}{5}/({0}{1}-{0}{2}/({0}{3}*{0}{4}))/{0}{4})'
                      .format(c, X['ticket'], X['tot_var'], X['tickets'],
                              X['dias'], X['fijos'])), C.DEC1)

    C.seccion(ws, 'A%d' % X['sec_exige'], 'LO QUE CADA ESCENARIO EXIGE')
    for letra in 'BCDE':
        motor.val(ws, letra + str(X['sec_exige']), '')
        ws[letra + str(X['sec_exige'])].fill = PatternFill(
            'solid', fgColor=C.CABECERA)
    linea(X['personal_ventas'], 'Coste de personal / ventas',
          lambda c: g(c, '%s$C$%d/%s%d' % (Q_PYG, P['personal'], c,
                                           X['ingresos'])), C.PCT,
          'La MISMA plantilla en los tres escenarios: cinco personas y tres '
          'jornadas y media. En el pesimista es donde se ve si esa plantilla '
          'se sostiene.')
    linea(X['ventas_jornada'], 'Ventas al año por jornada equivalente (€)',
          lambda c: g(c, '%s%d/%s$C$%d' % (c, X['ingresos'], Q_PER, R_TOT)),
          C.EUR,
          'Cuánto tendría que facturar cada jornada completa del cuadro de '
          'Personal. En el optimista dice si hace falta contratar antes de '
          'llegar ahí.')
    linea(X['saldo_est'],
          'Saldo de caja estimado al cierre del año 1 (mismo método en los tres)',
          lambda c: g(c, '{0}$B${1}+{2}{3}+{4}$C${5}-{6}$C${7}'
                      .format(Q_INV, I_FDM, c, X['neto'], Q_PYG,
                              P['amortizacion'], Q_FIN, F_ANIO_INI)), C.EUR,
          'Fondo de maniobra más el resultado del año, devolviendo la '
          'amortización (que no se paga) y restando el principal (que sí). Es '
          'una ESTIMACIÓN: el saldo de verdad, con su desfase de cobros y su '
          'liquidación de IVA, está en la hoja de Tesorería.')
    motor.val(ws, 'A%d' % X['saldo_real'],
              'Saldo real del mes 12 de la hoja de Tesorería (sólo el '
              'realista)')
    motor.f(ws, 'C%d' % X['saldo_real'],
            ie('IF(C%d="","",%sM%d)' % (X['ingresos'], Q_TES, T['saldo'])),
            fmt=C.EUR)
    C.nota(ws, 'E%d' % X['saldo_real'],
           'Sólo existe para el caso base: la tesorería mes a mes está '
           'construida sobre el año 1 del P&L, no sobre los escenarios.')
    C.parrafo(ws, X['nota'],
              'La columna «Realista» lee sus tres datos del año de crucero del '
              'P&L, así que reproduce EXACTAMENTE ese año. Si quieres mover el '
              'caso base, muévelo en «0. Supuestos» y esta columna te sigue.',
              'A', 'E', alto=30)
    ws.freeze_panes = 'A%d' % (X['cab'] + 1)   # hallazgo B12 (2026-09-10)
    C.pagina(ws, titulos='$%d:$%d' % (X['cab'], X['cab']),
             area='A1:E%d' % X['nota'])
    return ws


# ==========================================================================
# Hoja «Personal»
# ==========================================================================
def hoja_personal(wb):
    ws = wb.create_sheet(H_PER)
    C.anchos(ws, {'A': 34, 'B': 10, 'C': 10, 'D': 12, 'E': 15, 'F': 15,
                  'G': 15, 'H': 15, 'I': 15, 'J': 11, 'K': 62})
    C.encabezar(ws, 'Personal: cinco personas y tres jornadas y media',
                'Los cuatro perfiles son los del kit de tareas de pastelería, '
                'con DOS personas en Dependiente Vitrina. El bruto sale del '
                'convenio, en celda verde, y la Seguridad Social a cargo de la '
                'empresa va DENTRO: del bruto al coste empresa hay un tercio '
                'más.', col_fin='K')
    C.cabecera(ws, R_CAB,
               [('A', 'Persona y perfil'), ('B', 'Personas'), ('C', 'Jornada'),
                ('D', 'Grupo de convenio'), ('E', 'Bruto mes del grupo (€)'),
                ('F', 'Bruto mes del puesto (€)'),
                ('G', 'Seguridad Social a cargo de la empresa (€)'),
                ('H', 'Coste mes (€)'), ('I', 'Coste año (€)'),
                ('J', 'Horas/semana'), ('K', 'Notas')])

    for i, (pid, perfil, jornada, grupo, area, horas, turno) in \
            enumerate(D.PLANTILLA):
        fila = R_INI + i
        etiqueta = pid + ' · ' + perfil
        motor.val(ws, 'A%d' % fila, etiqueta)
        C.entrada(ws, 'B%d' % fila, 1, fmt=C.ENT, etiqueta=etiqueta,
                  align='center')
        C.entrada(ws, 'C%d' % fila, jornada, fmt=C.DEC1, etiqueta=etiqueta,
                  align='center')
        C.entrada(ws, 'D%d' % fila, 'Grupo %d' % grupo, etiqueta=etiqueta,
                  align='center')
        motor.f(ws, 'E%d' % fila,
                ie('INDEX($E$%d:$E$%d,MATCH(D%d,$A$%d:$A$%d,0))'
                   % (R_CONV_INI, R_CONV_FIN, fila, R_CONV_INI, R_CONV_FIN)),
                fmt=C.EUR)
        motor.f(ws, 'F%d' % fila, ie('E%d*C%d*B%d' % (fila, fila, fila)),
                fmt=C.EUR)
        motor.f(ws, 'G%d' % fila, ie('F%d*%s' % (fila, sup('ss'))), fmt=C.EUR)
        motor.f(ws, 'H%d' % fila, ie('F%d+G%d' % (fila, fila)), fmt=C.EUR)
        motor.f(ws, 'I%d' % fila, ie('H%d*MAX(12,%s)' % (fila, sup('pagas'))),
                fmt=C.EUR)
        C.entrada(ws, 'J%d' % fila, horas, fmt=C.ENT, etiqueta=etiqueta,
                  align='center')
        C.nota(ws, 'K%d' % fila, 'Área %s · turno de %s' % (area, turno))

    motor.val(ws, 'A%d' % R_TOT, 'TOTAL PLANTILLA', bold=True)
    for letra in ('B', 'C', 'F', 'G', 'H', 'I', 'J'):
        fmt = C.EUR if letra in 'FGHI' else (C.DEC1 if letra == 'C'
                                             else C.ENT)
        motor.f(ws, '%s%d' % (letra, R_TOT),
                ie('IF(COUNT(E%d:E%d)=0,"",SUM(%s%d:%s%d))'
                   % (R_INI, R_FIN, letra, R_INI, letra, R_FIN)), fmt=fmt,
                bold=True)
    C.total_fila(ws, R_TOT, 'ABCDEFGHIJK')
    C.nota(ws, 'K%d' % R_TOT,
           'Este «Coste año» es el que lee el P&L. No hay una segunda '
           'estimación de personal en ninguna parte del libro.')

    C.seccion(ws, 'A%d' % R_SEC_CONV,
              'EL CONVENIO DE REFERENCIA (sustitúyelo por el tuyo)')
    C.cabecera(ws, R_CONV_CAB,
               [('A', 'Grupo'), ('B', 'Denominación del grupo'),
                ('E', 'Bruto mes (€)'), ('F', 'Bruto año (€)'),
                ('G', 'Áreas funcionales'), ('I', 'Puestos que cita')],
               altura=26)
    ws.merge_cells('B%d:D%d' % (R_CONV_CAB, R_CONV_CAB))
    ws.merge_cells('G%d:H%d' % (R_CONV_CAB, R_CONV_CAB))
    ws.merge_cells('I%d:K%d' % (R_CONV_CAB, R_CONV_CAB))
    for i, (n, nombre, mes, anio, areas, puestos) in enumerate(D.CONVENIO):
        fila = R_CONV_INI + i
        motor.val(ws, 'A%d' % fila, 'Grupo %d' % n)
        ws.merge_cells('B%d:D%d' % (fila, fila))
        motor.val(ws, 'B%d' % fila, nombre, wrap=True)
        C.entrada(ws, 'E%d' % fila, mes, fmt=C.EUR,
                  etiqueta='Bruto mes del grupo %d' % n)
        motor.f(ws, 'F%d' % fila, ie('E%d*MAX(12,%s)' % (fila, sup('pagas'))),
                fmt=C.EUR)
        ws.merge_cells('G%d:H%d' % (fila, fila))
        motor.val(ws, 'G%d' % fila, ' / '.join(areas), wrap=True)
        ws.merge_cells('I%d:K%d' % (fila, fila))
        motor.val(ws, 'I%d' % fila, puestos or '-', wrap=True)
        ws.row_dimensions[fila].height = 26
    C.nota_celda(ws, 'E%d' % R_CONV_INI, 'PA-33',
                 'Convenio de Comercio e Industria de Confitería, Pastelería, '
                 'Bollería, Repostería, Heladería y Platos Cocinados de la '
                 'Comunidad de Madrid, código ' + D.CONVENIO_CODIGO
                 + ', revisión 2026 (' + D.CONVENIO_PUBLICACION + '). '
                 'Vigencia ' + D.CONVENIO_VIGENCIA + ': CADUCA el 31-12-2026.')
    C.dv_rango(ws, ['D%d' % (R_INI + i) for i in range(len(D.PLANTILLA))],
               "'%s'!$A$%d:$A$%d" % (H_PER, R_CONV_INI, R_CONV_FIN),
               'Elige un grupo del convenio',
               'Escribe uno de los grupos de la tabla del convenio de abajo.')

    C.seccion(ws, 'A%d' % R_SEC_SMI, 'CONTRASTE CON EL SMI')
    motor.val(ws, 'A%d' % R_SUELO, 'Bruto anual del grupo más bajo (€)')
    motor.f(ws, 'B%d' % R_SUELO,
            ie('IF(COUNT(F%d:F%d)=0,"",MIN(F%d:F%d))'
               % (R_CONV_INI, R_CONV_FIN, R_CONV_INI, R_CONV_FIN)), fmt=C.EUR)
    motor.val(ws, 'A%d' % R_SMI, 'SMI anual de referencia (€)')
    motor.f(ws, 'B%d' % R_SMI, ie(sup('smi')), fmt=C.EUR)
    C.nota_celda(ws, 'B%d' % R_SMI, 'PA-32')
    motor.val(ws, 'A%d' % R_VEREDICTO, '¿Algún grupo queda por debajo del SMI?')
    motor.f(ws, 'B%d' % R_VEREDICTO,
            ie('IF(B%d="","",IF(B%d<B%d,"Sí: revísalo con tu asesor",'
               '"No: en pastelería manda el convenio"))'
               % (R_SUELO, R_SUELO, R_SMI)))
    C.nota(ws, 'K%d' % R_VEREDICTO,
           'El SMI es una referencia ANUAL en cómputo global, no un mínimo por '
           'concepto. En pastelería manda el convenio: en Madrid el grupo más '
           'bajo ya lo supera. Las dos cosas caducan el 31-12-2026.')

    C.seccion(ws, 'A%d' % R_SEC_HORAS, 'HORAS Y DIMENSIONADO')
    motor.val(ws, 'A%d' % R_HORAS, 'Horas contratadas a la semana')
    motor.f(ws, 'B%d' % R_HORAS, ie('J%d' % R_TOT), fmt=C.ENT)
    motor.val(ws, 'A%d' % R_JORNADAS, 'Jornadas completas equivalentes')
    motor.f(ws, 'B%d' % R_JORNADAS, ie('C%d' % R_TOT), fmt=C.DEC1)
    motor.val(ws, 'A%d' % R_AVISOS,
              'Personas con más horas de las que da su jornada')
    motor.f(ws, 'B%d' % R_AVISOS,
            ie('SUMPRODUCT((J%d:J%d>C%d:C%d*%s)*1)'
               % (R_INI, R_FIN, R_INI, R_FIN, sup('horas_sem'))),
            fmt=C.ENT)
    motor.semaforo_isnumber(ws, 'B%d:B%d' % (R_AVISOS, R_AVISOS),
                            '$B$%d' % R_AVISOS, operador='>', umbral='0')
    C.nota(ws, 'K%d' % R_AVISOS,
           'Tiene que valer cero. Si se pone en rojo, alguien está contratado '
           'a media jornada y trabajando más horas de las que cobra: eso no es '
           'un ahorro, es una inspección.')
    C.parrafo(ws, R_NOTA,
              'El refuerzo de las campañas NO está en este cuadro: la '
              'estacionalidad manda tres personas más en Reyes y una o dos en '
              'el resto de los picos, y eso se dimensiona y se cuesta en los '
              'libros «estacionalidad-y-picos» y «plantilla-turnos-y-coste-'
              'personal». Aquí está la plantilla ESTABLE, que es la que entra '
              'en los costes fijos del P&L.', 'A', 'K', alto=44)
    ws.freeze_panes = 'A%d' % (R_CAB + 1)      # hallazgo B12 (2026-09-10)
    C.pagina(ws, titulos='$%d:$%d' % (R_CAB, R_CAB), area='A1:K%d' % R_NOTA)
    return ws


# ==========================================================================
# Hoja «Tesorería 12 meses»
# ==========================================================================
def hoja_tesoreria(wb):
    ws = wb.create_sheet(H_TES)
    C.anchos(ws, dict([('A', 50)] + [(c, 13) for c in MESES_COL]
                      + [('N', 15), ('O', 84)]))
    C.encabezar(ws, 'Tesorería de los 12 primeros meses',
                'El P&L dice si el negocio gana dinero; esta hoja dice si le '
                'queda caja para llegar a fin de mes. Es la primera que mira '
                'un banco, y la que más aperturas ha salvado.', col_fin='O')
    C.cabecera(ws, T['cab'],
               [('A', 'Concepto')]
               + [(c, 'Mes %d' % (i + 1)) for i, c in enumerate(MESES_COL)]
               + [('N', 'Año (€)'), ('O', 'Notas')])

    def sec(fila, texto):
        C.seccion(ws, 'A%d' % fila, texto)
        for letra in list(MESES_COL) + ['N', 'O']:
            motor.val(ws, letra + str(fila), '')
            ws[letra + str(fila)].fill = PatternFill('solid',
                                                     fgColor=C.CABECERA)

    # OJO: las filas de estacionalidad, rampa y reparto NO pueden llevar esta
    # guarda. «0. Supuestos» calcula desde ellas la actividad del año 1, y el
    # P&L calcula sus ingresos desde ese dato: guardarlas contra el P&L cierra
    # una referencia circular que Excel detecta y pycel no sabe resolver.
    G = 'IF({0}$B${1}="","",{2})'.format(Q_PYG, P['ingresos'], '%s')

    def g(expr):
        return ie(G % expr)

    # --- estacionalidad y rampa ------------------------------------------
    motor.val(ws, 'A%d' % T['estacionalidad'],
              'Estacionalidad del mes (tiene que sumar 100 %)')
    for i, col in enumerate(MESES_COL):
        C.entrada(ws, '%s%d' % (col, T['estacionalidad']),
                  round(D.estacionalidad_pesos()[i], 6), fmt=C.PCT2,
                  etiqueta='Estacionalidad del mes %d' % (i + 1))
    motor.f(ws, 'N%d' % T['estacionalidad'],
            ie('SUM(B%d:M%d)' % (T['estacionalidad'], T['estacionalidad'])),
            fmt=C.PCT)
    C.nota(ws, 'O%d' % T['estacionalidad'],
           'Enero manda por el roscón y diciembre por la Navidad; agosto es el '
           'suelo. Es el perfil de una ciudad media: en costa, agosto sube. '
           'Tiene que sumar 100 %, y la celda del año te lo dice.')

    motor.val(ws, 'A%d' % T['rampa'],
              'Rampa de arranque (% de la actividad de crucero)')
    for i, col in enumerate(MESES_COL):
        motor.f(ws, '%s%d' % (col, T['rampa']),
                ie('MIN(1,{0}+(1-{0})*{1}/MAX(1,{2}-1))'
                   .format(sup('rampa1'), i, sup('rampa_meses'))), fmt=C.PCT)
    C.nota(ws, 'O%d' % T['rampa'],
           'Sube en línea recta desde el porcentaje del mes 1 hasta el 100 % '
           'en los meses que digas en «0. Supuestos». Con el 100 % en el mes 1 '
           'la rampa desaparece, y entonces estás proyectando una apertura que '
           'no existe.')

    motor.val(ws, 'A%d' % T['reparto'],
              'Reparto de la actividad por mes (calculado)')
    for col in MESES_COL:
        motor.f(ws, '%s%d' % (col, T['reparto']),
                ie('{0}{1}*{0}{2}/SUMPRODUCT($B${1}:$M${1},$B${2}:$M${2})'
                   .format(col, T['estacionalidad'], T['rampa'])), fmt=C.PCT2)
    motor.f(ws, 'N%d' % T['reparto'],
            ie('SUM(B%d:M%d)' % (T['reparto'], T['reparto'])), fmt=C.PCT)
    C.nota(ws, 'O%d' % T['reparto'],
           'Estacionalidad x rampa, normalizado al 100 %: el AÑO factura lo que '
           'dice el P&L, pero repartido como lo reparte la realidad.')

    # --- cobros -----------------------------------------------------------
    sec(T['sec_cobros'], 'COBROS')
    motor.val(ws, 'A%d' % T['ventas'],
              'Ventas cobradas (con el IVA repercutido dentro)')
    for i, col in enumerate(MESES_COL):
        base = '{0}$B${1}*(1+{2})'.format(Q_PYG, P['ingresos'],
                                          sup('iva_medio'))
        if i == 0:
            reparto = '{0}{1}*(1-{2}/30)'.format(col, T['reparto'],
                                                 sup('dias_cobro'))
        else:
            ant = MESES_COL[i - 1]
            reparto = ('{0}{1}*(1-{3}/30)+{2}{1}*({3}/30)'
                       .format(col, T['reparto'], ant, sup('dias_cobro')))
        motor.f(ws, '%s%d' % (col, T['ventas']),
                g('%s*(%s)' % (base, reparto)), fmt=C.EUR)
    motor.f(ws, 'N%d' % T['ventas'],
            g('SUM(B%d:M%d)' % (T['ventas'], T['ventas'])), fmt=C.EUR)
    C.nota(ws, 'O%d' % T['ventas'],
           'Ventas del mes con el IVA repercutido, corregidas por los días '
           'medios de cobro que calcula la hoja de canales. El mostrador cobra '
           'al contado; el B2B, a 45 días.')

    # --- pagos ------------------------------------------------------------
    sec(T['sec_pagos'], 'PAGOS')
    motor.val(ws, 'A%d' % T['compras'],
              'Compras y coste de ventas (IVA incluido)')
    for i, col in enumerate(MESES_COL):
        base = '{0}$B${1}*(1+{0}$F${1})'.format(Q_PYG, P['coste_ventas'])
        if i == 0:
            reparto = '{0}{1}*(1-{2}/30)'.format(col, T['reparto'],
                                                 sup('dias_pago'))
        else:
            ant = MESES_COL[i - 1]
            reparto = ('{0}{1}*(1-{3}/30)+{2}{1}*({3}/30)'
                       .format(col, T['reparto'], ant, sup('dias_pago')))
        motor.f(ws, '%s%d' % (col, T['compras']),
                g('-%s*(%s)' % (base, reparto)), fmt=C.EUR)
    motor.f(ws, 'N%d' % T['compras'],
            g('SUM(B%d:M%d)' % (T['compras'], T['compras'])), fmt=C.EUR)
    C.nota(ws, 'O%d' % T['compras'],
           'Con los días medios de pago a proveedor de «0. Supuestos». A 30 '
           'días el mes 1 no paga nada: el proveedor te está financiando el '
           'arranque, y por eso perder ese crédito por un retraso sale tan '
           'caro.')

    motor.val(ws, 'A%d' % T['fijos'],
              'Costes fijos de explotación (IVA incluido)')
    for col in MESES_COL:
        motor.f(ws, '%s%d' % (col, T['fijos']),
                g('-(({0}$B${1}-{0}$B${2}-{0}$B${3}-{0}$B${4})+{0}$B${5})/12'
                  .format(Q_PYG, P['tot_fijos'], P['personal'],
                          P['amortizacion'], P['financieros'],
                          P['iva_fijos'])), fmt=C.EUR)
    motor.f(ws, 'N%d' % T['fijos'],
            g('SUM(B%d:M%d)' % (T['fijos'], T['fijos'])), fmt=C.EUR)
    C.nota(ws, 'O%d' % T['fijos'],
           'Los fijos que salen de caja, cada uno con su tipo de IVA: sin las '
           'nóminas (van abajo), sin la amortización (no se paga) y sin los '
           'intereses (van aparte).')

    motor.val(ws, 'A%d' % T['nominas'], 'Nóminas y Seguridad Social')
    for i, col in enumerate(MESES_COL):
        mes = i + 1
        extras = '+'.join(
            'IF(%d=%s,1,0)' % (mes, sup(k)) for k in ('extra1', 'extra2',
                                                      'extra3'))
        motor.f(ws, '%s%d' % (col, T['nominas']),
                g('-{0}$B${1}/MAX(12,{2})*(1+{3})'
                  .format(Q_PYG, P['personal'], sup('pagas'), extras)),
                fmt=C.EUR)
    motor.f(ws, 'N%d' % T['nominas'],
            g('SUM(B%d:M%d)' % (T['nominas'], T['nominas'])), fmt=C.EUR)
    C.nota(ws, 'O%d' % T['nominas'],
           'Doce mensualidades más las pagas extra en los meses que digas. Con '
           '15 pagas hay tres meses que salen dobles, y siempre pillan a '
           'contrapié: uno de ellos suele caer en el mes flojo.')

    motor.val(ws, 'A%d' % T['intereses'], 'Intereses del préstamo')
    motor.val(ws, 'A%d' % T['principal'], 'Devolución de principal del préstamo')
    for i, col in enumerate(MESES_COL):
        fila_mes = F_MES_INI + i
        for clave, letra in (('intereses', 'C'), ('principal', 'D')):
            motor.f(ws, '%s%d' % (col, T[clave]),
                    g('-IF({0}{1}{2}="",0,{0}{1}{2})'
                      .format(Q_FIN, letra, fila_mes)), fmt=C.EUR)
    for clave in ('intereses', 'principal'):
        motor.f(ws, 'N%d' % T[clave],
                g('SUM(B%d:M%d)' % (T[clave], T[clave])), fmt=C.EUR)
    C.nota(ws, 'O%d' % T['principal'],
           'Sale del cuadro MENSUAL de la hoja de Financiación. Durante la '
           'carencia esta fila vale cero, y ése es justo el efecto que se '
           'negocia con el banco.')

    # --- IVA --------------------------------------------------------------
    motor.val(ws, 'A%d' % T['iva_rep'], 'IVA repercutido del mes (memoria)')
    motor.val(ws, 'A%d' % T['iva_sop'], 'IVA soportado del mes (memoria)')
    for col in MESES_COL:
        motor.f(ws, '%s%d' % (col, T['iva_rep']),
                g('{0}$B${1}*{2}{3}*{4}'.format(Q_PYG, P['ingresos'], col,
                                                T['reparto'], sup('iva_medio'))),
                fmt=C.EUR)
        motor.f(ws, '%s%d' % (col, T['iva_sop']),
                g('{0}$B${1}*{2}{3}+{0}$B${4}/12'
                  .format(Q_PYG, P['iva_var'], col, T['reparto'],
                          P['iva_fijos'])), fmt=C.EUR)
    for clave in ('iva_rep', 'iva_sop'):
        motor.f(ws, 'N%d' % T[clave],
                g('SUM(B%d:M%d)' % (T[clave], T[clave])), fmt=C.EUR)

    motor.val(ws, 'A%d' % T['iva_arr'],
              'IVA a compensar arrastrado (inversión incluida)')
    motor.val(ws, 'A%d' % T['iva_liq'], 'Resultado de la liquidación trimestral')
    motor.val(ws, 'A%d' % T['iva_pago'], 'Pago del IVA (modelo 303)')
    TRIM = (('E', 'B', 'D', None), ('H', 'E', 'G', 'E'), ('K', 'H', 'J', 'H'))
    for col, ini, fin, previo in TRIM:
        if previo is None:
            motor.f(ws, '%s%d' % (col, T['iva_arr']),
                    g('%s$B$%d' % (Q_INV, I_IVA)), fmt=C.EUR)
        else:
            motor.f(ws, '%s%d' % (col, T['iva_arr']),
                    g('MAX(0,-%s%d)' % (previo, T['iva_liq'])), fmt=C.EUR)
        motor.f(ws, '%s%d' % (col, T['iva_liq']),
                g('SUM({0}{2}:{1}{2})-SUM({0}{3}:{1}{3})-{4}{5}'
                  .format(ini, fin, T['iva_rep'], T['iva_sop'], col,
                          T['iva_arr'])), fmt=C.EUR)
        motor.f(ws, '%s%d' % (col, T['iva_pago']),
                g('IF({0}{1}<=0,"",-{0}{1})'.format(col, T['iva_liq'])),
                fmt=C.EUR)
    motor.f(ws, 'N%d' % T['iva_pago'],
            g('IF(E{0}="",0,E{0})+IF(H{0}="",0,H{0})+IF(K{0}="",0,K{0})'
              .format(T['iva_pago'])), fmt=C.EUR)
    C.nota(ws, 'O%d' % T['iva_arr'],
           'El IVA soportado de la INVERSIÓN entra aquí y se compensa contra '
           'el repercutido de los trimestres siguientes. Con una inversión '
           'grande tardas medio año en recuperarlo, y ese medio año lo '
           'financias tú.')
    C.nota(ws, 'O%d' % T['iva_pago'],
           'Sólo sale caja cuando el resultado del trimestre es positivo. Los '
           'meses sin liquidación van en BLANCO, no a cero: cero significaría '
           '«liquidé y no pagué nada», que es otra cosa. El cuarto trimestre '
           'se liquida en enero del año siguiente y por eso no aparece.')

    # --- flujo y saldo ----------------------------------------------------
    motor.val(ws, 'A%d' % T['flujo'], 'FLUJO DEL MES', bold=True)
    for col in MESES_COL:
        motor.f(ws, '%s%d' % (col, T['flujo']),
                g('SUM({0}{1}:{0}{2})+IF({0}{3}="",0,{0}{3})'
                  .format(col, T['ventas'], T['principal'], T['iva_pago'])),
                fmt=C.EUR, bold=True)
    motor.f(ws, 'N%d' % T['flujo'],
            g('SUM(B%d:M%d)' % (T['flujo'], T['flujo'])), fmt=C.EUR, bold=True)
    C.total_fila(ws, T['flujo'], ['A'] + list(MESES_COL) + ['N', 'O'])

    motor.val(ws, 'A%d' % T['saldo'], 'SALDO ACUMULADO DE CAJA', bold=True)
    for i, col in enumerate(MESES_COL):
        if i == 0:
            expr = '%s$B$%d+%s%d' % (Q_INV, I_FDM, col, T['flujo'])
        else:
            expr = '%s%d+%s%d' % (MESES_COL[i - 1], T['saldo'], col,
                                  T['flujo'])
        motor.f(ws, '%s%d' % (col, T['saldo']), g(expr), fmt=C.EUR, bold=True)
    C.total_fila(ws, T['saldo'], ['A'] + list(MESES_COL) + ['N', 'O'])
    motor.semaforo_isnumber(ws, 'B%d:M%d' % (T['saldo'], T['saldo']),
                            'B%d' % T['saldo'], operador='<', umbral='0')

    motor.val(ws, 'A%d' % T['saldo_min'], 'Saldo mínimo del año (€)')
    motor.f(ws, 'B%d' % T['saldo_min'],
            g('IF(COUNT(B{0}:M{0})=0,"",MIN(B{0}:M{0}))'.format(T['saldo'])),
            fmt=C.EUR)
    motor.semaforo_isnumber(ws, 'B%d:B%d' % (T['saldo_min'], T['saldo_min']),
                            '$B$%d' % T['saldo_min'], operador='<', umbral='0')
    C.nota(ws, 'O%d' % T['saldo_min'],
           'Si es negativo te quedas sin caja: sube el fondo de maniobra, sube '
           'el préstamo o negocia más carencia. No lo dejes para verlo pasar.')
    motor.val(ws, 'A%d' % T['mes_fondo'], 'Mes en el que la caja toca fondo')
    motor.f(ws, 'B%d' % T['mes_fondo'],
            g('IF(COUNT(B{0}:M{0})=0,"",MATCH(MIN(B{0}:M{0}),B{0}:M{0},0))'
              .format(T['saldo'])), fmt=C.ENT)
    C.nota(ws, 'O%d' % T['mes_fondo'],
           'Es la respuesta a la pregunta que decide una operación bancaria: '
           'en qué mes se agota la caja.')

    # --- retorno ----------------------------------------------------------
    sec(T['sec_retorno'], 'RETORNO DE LA INVERSIÓN')
    for clave, col_pyg, anio in (('fcf1', 'B', 0), ('fcf2', 'C', 1),
                                 ('fcf3', 'D', 2)):
        motor.val(ws, 'A%d' % T[clave],
                  'Flujo de caja libre antes de la deuda, año %d' % (anio + 1))
        motor.f(ws, 'B%d' % T[clave],
                g('{0}{1}{2}+{0}{1}{3}+{4}$B${5}'
                  .format(Q_PYG, col_pyg, P['neto'], P['amortizacion'],
                          Q_FIN, F_ANIO_INI + anio)), fmt=C.EUR)
    C.nota(ws, 'O%d' % T['fcf1'],
           'Resultado neto más amortización (que no se paga) más los intereses '
           'de ese año: el principal no se resta porque este flujo se compara '
           'con la inversión, no con la deuda.')
    motor.val(ws, 'A%d' % T['inversion'],
              'Inversión a recuperar (sin el IVA, que se recupera por el 303)')
    motor.f(ws, 'B%d' % T['inversion'],
            g('%s$B$%d-%s$B$%d' % (Q_INV, I_NECES, Q_INV, I_IVA)), fmt=C.EUR)
    motor.val(ws, 'A%d' % T['payback'],
              'Payback del proyecto (años), antes de la deuda')
    motor.f(ws, 'B%d' % T['payback'],
            g('IF(B{1}>=B{0},ROUND(B{0}/B{1},1),'
              'IF(B{1}+B{2}>=B{0},ROUND(1+(B{0}-B{1})/B{2},1),'
              'IF(B{1}+B{2}+B{3}>=B{0},ROUND(2+(B{0}-B{1}-B{2})/B{3},1),'
              '"Más de 3 años")))'
              .format(T['inversion'], T['fcf1'], T['fcf2'], T['fcf3'])),
            fmt=C.DEC1)
    C.nota(ws, 'O%d' % T['payback'],
           'Cuánto tarda el negocio en devolver la inversión con lo que genera '
           'ANTES de pagar al banco. Si sale «Más de 3 años», este libro no '
           'proyecta lo suficiente para decírtelo: es una señal, no un fallo.')

    C.parrafo(ws, T['nota'],
              'Escribe los ingresos y los gastos SIN IVA en el resto del '
              'libro: la capa de IVA de esta hoja lo añade, porque la '
              'tesorería es caja. El Impuesto de Sociedades no está en este '
              'cuadro: se paga en julio del año siguiente, fuera de los doce '
              'meses.', 'A', 'O', alto=32)
    ws.freeze_panes = 'B%d' % (T['cab'] + 1)
    C.pagina(ws, titulos='$%d:$%d' % (T['cab'], T['cab']),
             area='A1:O%d' % T['nota'])
    return ws


# ==========================================================================
# Hoja «Financiación»
# ==========================================================================
def _cuadro_frances(principal, tipo_anual, plazo_meses, carencia_meses):
    """Mismo cuadro MENSUAL que construye la hoja de Financiación: interés sobre
    el saldo vivo, carencia de principal y cuota francesa algebraica. Devuelve
    la lista de intereses mes a mes."""
    i = tipo_anual / 12.0
    n_am = max(1, plazo_meses - carencia_meses)
    cuota = (principal / n_am if i == 0
             else principal * i / (1.0 - (1.0 + i) ** (-n_am)))
    saldo, intereses = principal, []
    for mes in range(1, plazo_meses + 1):
        interes = saldo * i
        intereses.append(interes)
        if mes > carencia_meses:
            saldo -= min(saldo, cuota - interes)
    return intereses


def _necesidad_estimada():
    """Necesidad total de caja al arranque, calculada en Python siguiendo la
    MISMA cadena que las hojas, para poder sembrar «Recursos propios» con un
    valor que deje la financiación cuadrada de fábrica.

    Existe para que la celda verde nazca con un número coherente (D21) y no con
    una cifra a ojo: `demo()` comprueba que la diferencia entre origen y usos
    del libro construido es menor de un euro.
    """
    importes, iva = [], 0.0
    for i, (bloque, _p, importe, _b, tipo, _f, _n) in enumerate(D.CAPEX):
        if bloque == 'Fondo de maniobra':
            continue
        v = DEFECTO_CAPEX[i] if importe is None else importe
        importes.append(v)
        iva += v * tipo
    subtotal = sum(importes)
    amortizable = (subtotal - D.CAPEX[7][2] - D.CAPEX[8][2])
    amort_anual = amortizable / float(D.ANIOS_AMORTIZACION)
    intereses = _cuadro_frances(D.FINANCIACION['principal'],
                                D.FINANCIACION['tipo_nominal'],
                                D.FINANCIACION['plazo_meses'],
                                D.FINANCIACION['carencia_meses'])
    int_a2 = sum(intereses[12:24])
    fijos = (D.coste_personal_anual() + 2400.0 * 12 + D.NEGOCIO['renta_mensual']
             * 12 + 1450.0 * 12 + sum(v for _c, _e, v, _t, _n in FIJOS_VERDES)
             + amort_anual + int_a2)
    fondo = D.P('meses_colchon_fondo_maniobra') * fijos / 12.0
    return subtotal + fondo + iva


def hoja_financiacion(wb):
    ws = wb.create_sheet(H_FIN)
    C.anchos(ws, {'A': 52, 'B': 16, 'C': 16, 'D': 18, 'E': 16, 'F': 17,
                  'G': 14, 'H': 74})
    C.encabezar(ws, 'Financiación: de dónde sale y cuánto cuesta devolverlo',
                'El cuadro de amortización es MENSUAL, con 84 filas, porque la '
                'carencia de este caso son seis meses y un cuadro anual no '
                'sabe expresar medio año. La cuota es una anualidad algebraica '
                'del sistema francés.', col_fin='H')

    def sec(fila, texto):
        C.seccion(ws, 'A%d' % fila, texto)
        for letra in 'BCDEFGH':
            motor.val(ws, letra + str(fila), '')
            ws[letra + str(fila)].fill = PatternFill('solid',
                                                     fgColor=C.CABECERA)

    def linea(fila, etiqueta, formula, fmt, nota_txt=None, bold=None,
              verde=None):
        motor.val(ws, 'A%d' % fila, etiqueta, bold=bold)
        if verde is not None:
            C.entrada(ws, 'B%d' % fila, verde, fmt=fmt, etiqueta=etiqueta,
                      bold=bold)
        else:
            motor.f(ws, 'B%d' % fila, formula, fmt=fmt, bold=bold)
        if nota_txt:
            C.nota(ws, 'H%d' % fila, nota_txt)

    sec(F['sec_origen'], 'ORIGEN DE FONDOS')
    linea(F['propios'], 'Recursos propios aportados', ie(sup('propios')),
          C.EUR, 'Aportación de los socios o del promotor. Un banco quiere ver '
          'un 25-30 % de fondos propios sobre la necesidad total antes de '
          'sentarse a hablar.')
    linea(F['prestamo'], 'Préstamo bancario', ie(sup('principal')), C.EUR)
    linea(F['otras'], 'Otras fuentes (ICO, ENISA, subvenciones, socios)', None,
          C.EUR, 'A cero por defecto. Las subvenciones suelen cobrarse DESPUÉS '
          'de justificar el gasto: no cuentes con ellas para pagar la obra.',
          verde=0.0)
    linea(F['tot_origen'], 'TOTAL ORIGEN DE FONDOS',
          ie('IF(COUNT(B%d:B%d)=0,"",SUM(B%d:B%d))'
             % (F['propios'], F['otras'], F['propios'], F['otras'])), C.EUR,
          bold=True)
    C.total_fila(ws, F['tot_origen'], 'ABCDEFGH')

    sec(F['sec_usos'], 'USOS')
    linea(F['necesidad'], 'Necesidad total de caja al arranque',
          ie('%s$B$%d' % (Q_INV, I_NECES)), C.EUR,
          'Inversión más el IVA que hay que adelantar.')
    linea(F['diferencia'], 'Diferencia (origen menos usos)',
          ie('B%d-B%d' % (F['tot_origen'], F['necesidad'])), C.EUR,
          'En negativo el plan NO está financiado. Muy por encima de cero '
          'tampoco es gratis: son intereses que pagas sin necesitarlos.')
    motor.semaforo_isnumber(ws, 'B%d:B%d' % (F['diferencia'], F['diferencia']),
                            '$B$%d' % F['diferencia'], operador='<', umbral='0')
    linea(F['dif_pct'], 'Diferencia sobre los usos (%)',
          ie('B%d/B%d' % (F['diferencia'], F['necesidad'])), C.PCT)
    linea(F['ajuste'], 'Préstamo que ajustaría el origen a la necesidad',
          ie('MAX(0,B%d-B%d+B%d)'
             % (F['necesidad'], F['tot_origen'], F['prestamo'])), C.EUR,
          'Cópialo en «Préstamo bancario solicitado» de «0. Supuestos» y la '
          'diferencia de arriba se pone a cero. Después mira el DSCR: pedir '
          'más siempre se puede, devolverlo no.')

    sec(F['sec_cond'], 'CONDICIONES DEL PRÉSTAMO')
    linea(F['principal'], 'Importe del principal',
          ie('IF(%s=0,"",%s)' % (sup('principal'), sup('principal'))), C.EUR)
    linea(F['tipo'], 'Tipo de interés nominal anual', ie(sup('tipo')), C.PCT2)
    linea(F['tipo_mes'], 'Tipo mensual', ie('B%d/12' % F['tipo']), '0.0000%')
    linea(F['plazo'], 'Plazo total (meses)', ie(sup('plazo')), C.ENT)
    linea(F['carencia'], 'Carencia de principal aplicada (meses)',
          ie('IF(%s>=%s,%s-1,%s)' % (sup('carencia'), sup('plazo'),
                                     sup('plazo'), sup('carencia'))),
          C.ENT,
          'Una carencia igual o mayor que el plazo no existe: la hoja la anula '
          'en origen.')
    linea(F['meses_am'], 'Meses de amortización',
          ie('B%d-B%d' % (F['plazo'], F['carencia'])), C.ENT)
    linea(F['cuota'], 'Cuota mensual durante la amortización',
          ie('IF(B{0}<=0,"",IF(B{1}=0,B{2}/B{0},B{2}*B{1}/(1-(1+B{1})^(-B{0}))))'
             .format(F['meses_am'], F['tipo_mes'], F['principal'])), C.EUR,
          bold=True)
    C.total_fila(ws, F['cuota'], 'ABCDEFGH')
    C.parrafo(ws, F['nota_cuota'],
              'Sistema francés, escrito como anualidad algebraica: capital x '
              'tipo mensual dividido por uno menos (uno más el tipo) elevado a '
              'menos el número de cuotas. Con el tipo al 0 % es el principal '
              'entre los meses de amortización. No se usa la función PMT: la '
              'familia la tiene prohibida porque no todos los motores de hoja '
              'de cálculo la evalúan igual.', 'A', 'H', alto=40)

    # --- cuadro mensual ---------------------------------------------------
    C.cabecera(ws, F['cab'],
               [('A', 'Mes'), ('B', 'Capital pendiente (€)'),
                ('C', 'Intereses (€)'), ('D', 'Amortización de principal (€)'),
                ('E', 'Cuota total (€)'), ('F', 'Capital al cierre (€)'),
                ('G', ''), ('H', 'Notas')])
    for i in range(F_MESES):
        fila = F_MES_INI + i
        motor.val(ws, 'A%d' % fila, i + 1, fmt=C.ENT, align='center')
        if i == 0:
            motor.f(ws, 'B%d' % fila,
                    ie('IF($B$%d="","",$B$%d)' % (F['principal'],
                                                  F['principal'])), fmt=C.EUR)
        else:
            motor.f(ws, 'B%d' % fila,
                    ie('IF(OR(A{0}>$B${1},F{2}=""),"",F{2})'
                       .format(fila, F['plazo'], fila - 1)), fmt=C.EUR)
        motor.f(ws, 'C%d' % fila,
                ie('IF(OR(A{0}>$B${1},B{0}=""),"",B{0}*$B${2})'
                   .format(fila, F['plazo'], F['tipo_mes'])), fmt=C.EUR)
        motor.f(ws, 'D%d' % fila,
                ie('IF(OR(A{0}>$B${1},B{0}=""),"",IF(A{0}<=$B${2},0,'
                   'MIN(B{0},$B${3}-C{0})))'
                   .format(fila, F['plazo'], F['carencia'], F['cuota'])),
                fmt=C.EUR)
        motor.f(ws, 'E%d' % fila,
                ie('IF(C{0}="","",C{0}+D{0})'.format(fila)), fmt=C.EUR)
        motor.f(ws, 'F%d' % fila,
                ie('IF(D{0}="","",B{0}-D{0})'.format(fila)), fmt=C.EUR)
    C.nota(ws, 'H%d' % F_MES_INI,
           'Durante la carencia la columna de amortización vale cero y sólo se '
           'pagan intereses. Es lo que hace que el préstamo no se coma la caja '
           'en la rampa de arranque, y es lo primero que hay que negociar.')

    motor.val(ws, 'A%d' % F_TOL, 'Tolerancia de cierre del cuadro (€)')
    C.entrada(ws, 'B%d' % F_TOL, 0.01, fmt=C.EUR,
              etiqueta='Tolerancia de cierre del cuadro')
    motor.val(ws, 'A%d' % F_PENDIENTE, 'Capital pendiente al vencimiento')
    motor.f(ws, 'B%d' % F_PENDIENTE,
            ie('IF($B${0}="","",ROUND($B${0}-SUM(D{1}:D{2}),2))'
               .format(F['principal'], F_MES_INI, F_MES_FIN)), fmt=C.EUR)
    motor.val(ws, 'A%d' % F_CIERRA, '¿Cierra el cuadro de amortización?')
    motor.f(ws, 'B%d' % F_CIERRA,
            ie('IF(B{0}="","",IF(ABS(B{0})<=B{1},"Sí","No: revisa el plazo"))'
               .format(F_PENDIENTE, F_TOL)))
    C.nota(ws, 'H%d' % F_CIERRA,
           'Tiene que decir «Sí»: el principal devuelto en el cuadro tiene que '
           'ser el prestado. Si dice que no, el plazo no da para amortizarlo '
           'con esa cuota y queda capital sin devolver al vencimiento.')

    # --- resumen por año --------------------------------------------------
    C.seccion(ws, 'A%d' % F_SEC_ANIO, 'RESUMEN POR AÑO Y COBERTURA DEL '
                                      'SERVICIO DE LA DEUDA (DSCR)')
    C.cabecera(ws, F_ANIO_CAB,
               [('A', 'Año'), ('B', 'Intereses (€)'),
                ('C', 'Principal devuelto (€)'), ('D', 'Cuota total (€)'),
                ('E', 'Flujo disponible para la deuda (€)'), ('F', 'DSCR'),
                ('G', ''), ('H', 'Notas')], altura=30)
    for k in range(F_ANIOS):
        fila = F_ANIO_INI + k
        ini = F_MES_INI + k * 12
        fin = ini + 11
        motor.val(ws, 'A%d' % fila, k + 1, fmt=C.ENT, align='center')
        motor.f(ws, 'B%d' % fila, ie('SUM(C%d:C%d)' % (ini, fin)), fmt=C.EUR)
        motor.f(ws, 'C%d' % fila, ie('SUM(D%d:D%d)' % (ini, fin)), fmt=C.EUR)
        motor.f(ws, 'D%d' % fila, ie('B%d+C%d' % (fila, fila)), fmt=C.EUR)
        col_pyg = 'BCD'[k] if k < 3 else 'D'
        motor.f(ws, 'E%d' % fila,
                ie('{0}{1}{2}-{0}{1}{3}+{0}{1}{4}+B{5}'
                   .format(Q_PYG, col_pyg, P['rai'], P['is'],
                           P['amortizacion'], fila)), fmt=C.EUR)
        motor.f(ws, 'F%d' % fila,
                ie('IF(OR(D{0}="",D{0}=0),"",E{0}/D{0})'.format(fila)),
                fmt=C.DEC)
    C.nota(ws, 'H%d' % F_ANIO_INI,
           'El flujo disponible se calcula ANTES de la deuda y DESPUÉS de '
           'impuestos: al resultado antes de impuestos se le resta la cuota '
           'del Impuesto de Sociedades y se le suma la amortización, que no se '
           'paga. Del año 4 en adelante se mantiene el flujo del año 3: este '
           'libro no proyecta más lejos, y decirlo es más honesto que '
           'inventarse un crecimiento.')
    motor.val(ws, 'A%d' % F_DSCR_MIN, 'DSCR mínimo de todo el cuadro',
              bold=True)
    motor.f(ws, 'B%d' % F_DSCR_MIN,
            ie('IF(COUNT(F{0}:F{1})=0,"",MIN(F{0}:F{1}))'
               .format(F_ANIO_INI, F_ANIO_FIN)), fmt=C.DEC, bold=True)
    C.total_fila(ws, F_DSCR_MIN, 'ABCDEFGH')
    motor.val(ws, 'A%d' % F_DSCR_SEM, 'Lectura del DSCR')
    motor.f(ws, 'B%d' % F_DSCR_SEM,
            ie('IF(B{0}="","",IF(B{0}<{1},"Por debajo del mínimo",'
               'IF(B{0}<{2},"Entre el mínimo y el objetivo","En objetivo")))'
               .format(F_DSCR_MIN, sup('dscr_min'), sup('dscr_obj'))))
    motor.semaforo_isnumber(ws, 'B%d:B%d' % (F_DSCR_MIN, F_DSCR_MIN),
                            '$B$%d' % F_DSCR_MIN, operador='<',
                            umbral=sup('dscr_min'))
    C.nota(ws, 'H%d' % F_DSCR_MIN,
           'Es el PEOR año de los que dura el préstamo, no el mejor de los '
           'tres que proyecta el P&L. Es el número que decide la operación.')
    C.parrafo(ws, F_NOTA,
              'El año de carencia infla el DSCR: con sólo intereses en el '
              'denominador el ratio se dispara y no está midiendo el negocio, '
              'está midiendo que todavía no devuelves nada. Mira el año 2 y el '
              'año 3, que es cuando la cuota está completa.', 'A', 'H',
              alto=32)
    ws.freeze_panes = 'A%d' % (F['cab'] + 1)
    C.pagina(ws, titulos='$%d:$%d' % (F['cab'], F['cab']),
             area='A1:H%d' % F_NOTA)
    return ws


# ==========================================================================
# Hoja «Canales y Punto Muerto» (absorbida, D3)
# ==========================================================================
def hoja_canales(wb):
    ws = wb.create_sheet(H_CAN)
    C.anchos(ws, {'A': 44, 'B': 12, 'C': 12, 'D': 13, 'E': 14, 'F': 11,
                  'G': 15, 'H': 16, 'I': 17, 'J': 78})
    C.encabezar(ws, 'Canales y punto muerto: qué canal sostiene el negocio',
                'Cada canal tiene su margen, su coste de servir, su comisión y '
                'sus días de cobro. La suma no es lo interesante: lo '
                'interesante es qué pasa con el punto muerto cuando quitas '
                'uno.', col_fin='J')
    C.cabecera(ws, K_CAB,
               [('A', 'Canal'), ('B', '% de las ventas'),
                ('C', 'Margen bruto (%)'), ('D', 'Coste de servir (%)'),
                ('E', 'Comisión de plataforma (%)'), ('F', 'Días de cobro'),
                ('G', 'Margen de contribución (%)'),
                ('H', 'Ventas del año de crucero (€)'),
                ('I', 'Margen de contribución (€)'), ('J', 'Notas')])

    for i, c in enumerate(D.CANALES):
        fila = K_INI + i
        motor.val(ws, 'A%d' % fila, c['canal'], wrap=True)
        C.entrada(ws, 'B%d' % fila, round(c['ventas_pct'] / 100.0, 4),
                  fmt=C.PCT, etiqueta=c['canal'] + ' · % de ventas')
        C.entrada(ws, 'C%d' % fila, c['margen_bruto'], fmt=C.PCT,
                  etiqueta=c['canal'] + ' · margen bruto')
        C.entrada(ws, 'D%d' % fila, c['coste_servir_pct'], fmt=C.PCT,
                  etiqueta=c['canal'] + ' · coste de servir')
        C.entrada(ws, 'E%d' % fila, c['comision_plataforma'], fmt=C.PCT,
                  etiqueta=c['canal'] + ' · comisión de plataforma')
        C.entrada(ws, 'F%d' % fila, c['cobro_dias'], fmt=C.ENT,
                  etiqueta=c['canal'] + ' · días de cobro', align='center')
        motor.f(ws, 'G%d' % fila,
                ie('C{0}-D{0}-E{0}'.format(fila)), fmt=C.PCT)
        motor.f(ws, 'H%d' % fila,
                ie('IF({0}$C${1}="","",{0}$C${1}*B{2})'
                   .format(Q_PYG, P['ingresos'], fila)), fmt=C.EUR)
        motor.f(ws, 'I%d' % fila, ie('IF(H{0}="","",H{0}*G{0})'.format(fila)),
                fmt=C.EUR)
        C.nota(ws, 'J%d' % fila, c['nota'])
        ws.row_dimensions[fila].height = 44

    motor.val(ws, 'A%d' % K_TOT, 'TOTAL', bold=True)
    for letra in ('B', 'H', 'I'):
        motor.f(ws, '%s%d' % (letra, K_TOT),
                ie('IF(COUNT({0}{1}:{0}{2})=0,"",SUM({0}{1}:{0}{2}))'
                   .format(letra, K_INI, K_FIN)),
                fmt=(C.PCT if letra == 'B' else C.EUR), bold=True)
    motor.f(ws, 'G%d' % K_TOT, ie('I%d/H%d' % (K_TOT, K_TOT)), fmt=C.PCT,
            bold=True)
    C.total_fila(ws, K_TOT, 'ABCDEFGHIJ')
    C.nota(ws, 'J%d' % K_TOT,
           'La columna «% de las ventas» tiene que sumar 100 %. Si no suma, '
           'todo lo de abajo está mal.')

    def linea(fila, etiqueta, formula, fmt=None, nota_txt=None, bold=None,
              verde=None, id_pa=None):
        motor.val(ws, 'A%d' % fila, etiqueta, bold=bold)
        if verde is not None:
            C.entrada(ws, 'B%d' % fila, verde, fmt=fmt, etiqueta=etiqueta)
        else:
            motor.f(ws, 'B%d' % fila, formula, fmt=fmt, bold=bold)
        if id_pa:
            C.nota_celda(ws, 'B%d' % fila, id_pa)
        if nota_txt:
            C.nota(ws, 'J%d' % fila, nota_txt)

    C.seccion(ws, 'A%d' % K_SEC_PM, 'EL PUNTO MUERTO, CANAL A CANAL')
    linea(K_FIJOS, 'Costes fijos del año de crucero (€)',
          ie('%s$C$%d' % (Q_PYG, P['tot_fijos'])), C.EUR)
    linea(K_FIJOS_MES, 'Costes fijos mensuales (€)',
          ie('B%d/12' % K_FIJOS), C.EUR)
    linea(K_MC, 'Margen de contribución medio ponderado (%)',
          ie('G%d' % K_TOT), C.PCT)
    linea(K_REGLA, 'Margen bruto objetivo de la regla única (%)',
          ie(sup('margen_obj')), C.PCT)
    linea(K_DIF_MC, 'Lo que se llevan servir, cobrar y las comisiones (%)',
          ie('B%d-B%d' % (K_REGLA, K_MC)), C.PCT,
          'La diferencia entre el margen bruto de la carta y el margen de '
          'contribución real. No es un error de ninguno de los dos: es lo que '
          'cuesta poner el producto en manos del cliente por cada vía.')
    linea(K_PM_CON, 'Punto muerto mensual CON todos los canales (€)',
          ie('IF(B{0}<=0,"",B{1}/B{0})'.format(K_MC, K_FIJOS_MES)), C.EUR,
          bold=True)
    linea(K_PM_SIN, 'Punto muerto mensual SIN el canal B2B (€)',
          ie('IF(OR(H{0}-H{1}<=0,(I{0}-I{1})/(H{0}-H{1})<=0),"",'
             'B{2}/((I{0}-I{1})/(H{0}-H{1})))'
             .format(K_TOT, K_B2B, K_FIJOS_MES)), C.EUR, bold=True)
    linea(K_PM_DIF, 'Diferencia (€/mes)',
          ie('IF(OR(B{0}="",B{1}=""),"",B{0}-B{1})'
             .format(K_PM_SIN, K_PM_CON)), C.EUR,
          'Si sale NEGATIVA, quitar el B2B te BAJA el punto muerto: ese canal '
          'estaba tirando del margen medio hacia abajo, y lo que aporta es '
          'volumen, no rentabilidad. Ojo, eso no dice que haya que cerrarlo: '
          'dice qué tienes que exigirle.')
    C.total_fila(ws, K_PM_CON, 'ABCDEFGHIJ')
    C.total_fila(ws, K_PM_SIN, 'ABCDEFGHIJ')
    linea(K_SOSTIENE, 'Canal que sostiene el negocio',
          ie('IF(COUNT(I{0}:I{1})=0,"",INDEX($A${0}:$A${1},'
             'MATCH(MAX($I${0}:$I${1}),$I${0}:$I${1},0)))'
             .format(K_INI, K_FIN)), None, bold=True)
    linea(K_APORTE, 'Aporte de ese canal al margen de contribución (%)',
          ie('IF(I{0}="","",MAX($I${1}:$I${2})/I{0})'
             .format(K_TOT, K_INI, K_FIN)), C.PCT,
          'Es el canal que paga el alquiler. Todo lo que pongas en marcha en '
          'los otros tres se financia con éste, así que lo primero que hay que '
          'proteger es su margen.')

    C.seccion(ws, 'A%d' % K_SEC_ART3,
              'EL AVISO DEL ART. 3: CUÁNDO EL B2B TE CAMBIA DE RÉGIMEN')
    linea(K_PESO_B2B, 'Peso del B2B sobre las ventas (%)',
          ie('B%d' % K_B2B), C.PCT)
    linea(K_UMBRAL, 'Umbral de la vía a) del art. 3.2 (%)', None, C.PCT,
          'Una actividad es marginal si el suministro a otros minoristas es '
          'menor o igual al 25 % del volumen anual O si la comercialización '
          'total no pasa de 500 kg a la semana. Son vías ALTERNATIVAS: basta '
          'cumplir una. La de los kilos se comprueba en el libro '
          '«checklist-legal-y-licencias».', verde=0.25, id_pa='PA-02b')
    linea(K_VEREDICTO, 'Lectura del art. 3 por la vía del porcentaje',
          ie('IF(B{0}=0,"Sin B2B: el art. 3 no te aplica",'
             'IF(B{0}<={1},"Marginal por la vía a): sigue la comprobación en '
             'el libro de licencias","Por encima del 25 %: comprueba la vía de '
             'los 500 kg antes de nada"))'.format(K_PESO_B2B, 'B%d' % K_UMBRAL)),
          None, bold=True)
    C.nota_celda(ws, 'B%d' % K_VEREDICTO, 'PA-02')
    C.parrafo(ws, K_VEREDICTO + 1,
              'El art. 3 del RD 1021/2022 sólo se activa si suministras a '
              'establecimientos minoristas de DISTINTA titularidad. Si entras, '
              'las tres condiciones son ACUMULATIVAS -marginal Y localizado Y '
              'restringido- y hay que presentar declaración responsable y '
              'llevar registro de destinatarios, cantidades y fechas. Este '
              'aviso mira sólo el porcentaje: el árbol completo está en el '
              'libro «checklist-legal-y-licencias».', 'A', 'J', alto=44)

    linea(K_COBRO, 'Días medios de cobro ponderados',
          ie(sup('dias_cobro')), C.DEC1,
          'Es el número que usa la hoja de Tesorería para desfasar los cobros. '
          'Sube solo en cuanto crece el B2B, y ése es su coste oculto: no es '
          'sólo que deje menos margen, es que lo deja más tarde.')
    C.parrafo(ws, K_NOTA,
              'Los porcentajes de esta hoja son SUPUESTOS salvo el peso del '
              'B2B, que sale de la única cifra publicada de estructura de '
              'cobro de una pastelería española: 95 % de venta directa y 5 % a '
              'crédito a 30-60 días. El reparto entre mostrador, encargos y '
              'online es una hipótesis de trabajo: cámbiala por la tuya en '
              'cuanto tengas tres meses de datos reales del TPV.', 'A', 'J',
              alto=44)
    ws.freeze_panes = 'A%d' % (K_CAB + 1)      # hallazgo B12 (2026-09-10)
    C.pagina(ws, titulos='$%d:$%d' % (K_CAB, K_CAB), area='A1:J%d' % K_NOTA)
    return ws


# ==========================================================================
# Mapa de celdas citables
# ==========================================================================
def mapa():
    m = [
        # --- supuestos ---------------------------------------------------
        ('Tickets al día en velocidad de crucero', H_SUP,
         'B%d' % S['tickets'], 'entrada'),
        ('Piezas por ticket', H_SUP, 'B%d' % S['piezas'], 'entrada'),
        ('Ticket medio con IVA', H_SUP, 'B%d' % S['ticket_iva'], 'salida'),
        ('Ticket medio sin IVA', H_SUP, 'B%d' % S['ticket_sin'], 'salida'),
        ('Días de apertura al año', H_SUP, 'B%d' % S['dias'], 'entrada'),
        ('Actividad del año 1 sobre la de crucero', H_SUP,
         'B%d' % S['factor_a1'], 'salida'),
        ('IVA medio ponderado de la carta', H_SUP, 'B%d' % S['iva_medio'],
         'salida'),
        ('Food cost objetivo del plan (regla única)', H_SUP,
         'B%d' % S['food_cost'], 'parametro'),
        ('Margen bruto objetivo', H_SUP, 'B%d' % S['margen_obj'], 'salida'),
        ('Seguridad Social a cargo de la empresa', H_SUP, 'B%d' % S['ss'],
         'parametro'),
        ('Pagas del convenio', H_SUP, 'B%d' % S['pagas'], 'parametro'),
        ('Alquiler mensual del local', H_SUP, 'B%d' % S['renta'], 'entrada'),
        ('Retribución mensual del propietario', H_SUP,
         'B%d' % S['retribucion'], 'entrada'),
        ('Recursos propios aportados', H_SUP, 'B%d' % S['propios'], 'entrada'),
        ('Principal del préstamo', H_SUP, 'B%d' % S['principal'], 'entrada'),
        ('Tipo de interés nominal anual', H_SUP, 'B%d' % S['tipo'], 'entrada'),
        ('Meses de fondo de maniobra', H_SUP, 'B%d' % S['colchon'], 'entrada'),
        ('Margen neto de referencia, suelo', H_SUP, 'B%d' % S['neto_suelo'],
         'parametro'),
        ('Margen neto de referencia, techo', H_SUP, 'B%d' % S['neto_techo'],
         'parametro'),
        # --- inversión ----------------------------------------------------
        ('Obra civil e instalaciones', H_INV, 'B%d' % _FILA_CAPEX[0],
         'entrada'),
        ('Equipamiento de obrador', H_INV, 'B%d' % _FILA_CAPEX[1], 'entrada'),
        ('Equipamiento de tienda y vitrina', H_INV, 'B%d' % _FILA_CAPEX[2],
         'entrada'),
        ('Subtotal de inversión sin el fondo de maniobra', H_INV,
         'B%d' % I_SUBTOT, 'salida'),
        ('Fondo de maniobra', H_INV, 'B%d' % I_FDM, 'salida'),
        ('Inversión total', H_INV, 'B%d' % I_TOTAL, 'salida'),
        ('CAPEX sin el fondo de maniobra', H_INV, 'B%d' % I_CAPEX, 'salida'),
        ('IVA soportado sobre la inversión', H_INV, 'B%d' % I_IVA, 'salida'),
        ('Necesidad total de caja al arranque', H_INV, 'B%d' % I_NECES,
         'salida'),
        ('Inmovilizado amortizable', H_INV, 'B%d' % I_BASE, 'salida'),
        ('Amortización anual del inmovilizado', H_INV, 'B%d' % I_AMORT,
         'salida'),
        # --- P&L -----------------------------------------------------------
        ('Ingresos del año 1', H_PYG, 'B%d' % P['ingresos'], 'salida'),
        ('Ingresos del año de crucero', H_PYG, 'C%d' % P['ingresos'],
         'salida'),
        ('Ingresos del año 3', H_PYG, 'D%d' % P['ingresos'], 'salida'),
        ('Coste de ventas del año de crucero', H_PYG,
         'C%d' % P['coste_ventas'], 'salida'),
        ('Margen bruto del año de crucero', H_PYG, 'C%d' % P['margen'],
         'salida'),
        ('Coste de personal anual', H_PYG, 'C%d' % P['personal'], 'salida'),
        ('Retribución anual del propietario', H_PYG, 'C%d' % P['propietario'],
         'salida'),
        ('Total de costes fijos del año 1', H_PYG, 'B%d' % P['tot_fijos'],
         'salida'),
        ('Total de costes fijos del año de crucero', H_PYG,
         'C%d' % P['tot_fijos'], 'salida'),
        ('Resultado antes de impuestos del año 1', H_PYG, 'B%d' % P['rai'],
         'salida'),
        ('Resultado neto del año 1', H_PYG, 'B%d' % P['neto'], 'salida'),
        ('Resultado neto del año de crucero', H_PYG, 'C%d' % P['neto'],
         'salida'),
        ('Resultado neto del año 3', H_PYG, 'D%d' % P['neto'], 'salida'),
        ('Margen neto del año de crucero', H_PYG, 'C%d' % P['r_neto'],
         'salida'),
        ('Margen neto del año de crucero dentro de la banda 8-12 %', H_PYG,
         'C%d' % P['r_banda'], 'salida'),
        ('Coste de personal sobre ventas en el año de crucero', H_PYG,
         'C%d' % P['r_personal'], 'salida'),
        # --- punto de equilibrio -------------------------------------------
        ('Margen de contribución por ticket', H_PEQ, 'C%d' % E['mc_ticket'],
         'salida'),
        ('Tickets al día para el equilibrio contable', H_PEQ,
         'C%d' % E['tk_dia'], 'salida'),
        ('Ingresos mensuales de equilibrio contable', H_PEQ,
         'C%d' % E['ing_mes'], 'salida'),
        ('Tickets al día para el equilibrio de caja', H_PEQ,
         'C%d' % E['tkc_dia'], 'salida'),
        ('Holgura sobre el equilibrio de caja en el año de crucero', H_PEQ,
         'C%d' % E['holgura'], 'salida'),
        # --- escenarios ------------------------------------------------------
        ('Ingresos del escenario pesimista', H_ESC, 'B%d' % X['ingresos'],
         'salida'),
        ('Resultado neto del escenario pesimista', H_ESC, 'B%d' % X['neto'],
         'salida'),
        ('Resultado neto del escenario optimista', H_ESC, 'D%d' % X['neto'],
         'salida'),
        ('Coste de personal sobre ventas en el escenario pesimista', H_ESC,
         'B%d' % X['personal_ventas'], 'salida'),
        # --- personal ---------------------------------------------------------
        ('Coste anual del Jefe Pastelero', H_PER, 'I%d' % R_INI, 'salida'),
        ('Coste mensual total de la plantilla', H_PER, 'H%d' % R_TOT,
         'salida'),
        ('Coste anual total de la plantilla', H_PER, 'I%d' % R_TOT, 'salida'),
        ('Jornadas completas equivalentes', H_PER, 'C%d' % R_TOT, 'salida'),
        ('Horas contratadas a la semana', H_PER, 'J%d' % R_TOT, 'salida'),
        ('Bruto anual del grupo más bajo del convenio', H_PER,
         'B%d' % R_SUELO, 'salida'),
        ('Veredicto del contraste con el SMI', H_PER, 'B%d' % R_VEREDICTO,
         'salida'),
        # --- tesorería --------------------------------------------------------
        ('Saldo de caja del mes 12', H_TES, 'M%d' % T['saldo'], 'salida'),
        ('Saldo mínimo de caja del año 1', H_TES, 'B%d' % T['saldo_min'],
         'salida'),
        ('Mes en el que la caja toca fondo', H_TES, 'B%d' % T['mes_fondo'],
         'salida'),
        ('Flujo de caja libre del año 1 antes de la deuda', H_TES,
         'B%d' % T['fcf1'], 'salida'),
        ('Inversión a recuperar', H_TES, 'B%d' % T['inversion'], 'salida'),
        ('Payback del proyecto en años', H_TES, 'B%d' % T['payback'],
         'salida'),
        # --- financiación -----------------------------------------------------
        ('Total del origen de fondos', H_FIN, 'B%d' % F['tot_origen'],
         'salida'),
        ('Diferencia entre origen y usos', H_FIN, 'B%d' % F['diferencia'],
         'salida'),
        ('Cuota mensual del préstamo', H_FIN, 'B%d' % F['cuota'], 'salida'),
        ('Intereses del año 1', H_FIN, 'B%d' % F_ANIO_INI, 'salida'),
        ('Principal devuelto en el año 1', H_FIN, 'C%d' % F_ANIO_INI,
         'salida'),
        ('Intereses del año 2', H_FIN, 'B%d' % (F_ANIO_INI + 1), 'salida'),
        ('Capital pendiente al vencimiento', H_FIN, 'B%d' % F_PENDIENTE,
         'salida'),
        ('El cuadro de amortización cierra', H_FIN, 'B%d' % F_CIERRA,
         'salida'),
        ('DSCR del año 2', H_FIN, 'F%d' % (F_ANIO_INI + 1), 'salida'),
        ('DSCR mínimo de todo el cuadro', H_FIN, 'B%d' % F_DSCR_MIN, 'salida'),
        ('Lectura del DSCR', H_FIN, 'B%d' % F_DSCR_SEM, 'salida'),
        # --- canales -----------------------------------------------------------
        ('Margen de contribución del mostrador', H_CAN, 'G%d' % K_INI,
         'salida'),
        ('Margen de contribución del canal B2B', H_CAN, 'G%d' % K_B2B,
         'salida'),
        ('Ventas anuales del mostrador', H_CAN, 'H%d' % K_INI, 'salida'),
        ('Margen de contribución medio ponderado', H_CAN, 'B%d' % K_MC,
         'salida'),
        ('Lo que se llevan servir, cobrar y las comisiones', H_CAN,
         'B%d' % K_DIF_MC, 'salida'),
        ('Punto muerto mensual con todos los canales', H_CAN,
         'B%d' % K_PM_CON, 'salida'),
        ('Punto muerto mensual sin el canal B2B', H_CAN, 'B%d' % K_PM_SIN,
         'salida'),
        ('Diferencia del punto muerto al quitar el B2B', H_CAN,
         'B%d' % K_PM_DIF, 'salida'),
        ('Canal que sostiene el negocio', H_CAN, 'B%d' % K_SOSTIENE, 'salida'),
        ('Aporte del canal que sostiene el negocio', H_CAN, 'B%d' % K_APORTE,
         'salida'),
        ('Peso del B2B sobre las ventas', H_CAN, 'B%d' % K_PESO_B2B, 'salida'),
        ('Lectura del art. 3 por la vía del porcentaje', H_CAN,
         'B%d' % K_VEREDICTO, 'salida'),
        ('Días medios de cobro ponderados', H_CAN, 'B%d' % K_COBRO, 'salida'),
    ]
    return m


NOTAS_MAPA = (
    'Libro HÍBRIDO: molde de planes (motor 2.2) más la hoja propia «Canales y '
    'Punto Muerto», absorbida por la decisión D3. La versión que manda es la '
    'de la familia de guías (1.0 · septiembre 2026). El AÑO 2 del P&L es el '
    'año de CRUCERO y reproduce `datos_ejemplo.cuenta_resultados_crucero()`; '
    'el año 1 lleva la rampa de arranque. Los intereses del año 1 se calculan '
    'con el cuadro francés mensual y difieren en unos pocos euros de la '
    'aproximación lineal de `datos_ejemplo.intereses_anio_1()`: manda este '
    'libro. Cita SIEMPRE la columna del año que quieras nombrar.'
)


#: Las TRES fórmulas del libro cuyo resultado correcto es la cadena vacía. Son
#: los pagos trimestrales del modelo 303: mientras el IVA soportado de la
#: inversión siga arrastrándose, la liquidación sale negativa y NO se paga
#: nada. La celda va en blanco a propósito, porque un 0 significaría «liquidé y
#: no salió a pagar», que es otra cosa. `demo()` comprueba con pycel que valen
#: `''` y no un error.
BLANCOS_OK = dict(
    (("%s" % H_TES, "%s%d" % (col, T['iva_pago'])),
     'Liquidación del trimestre negativa: no se paga, se compensa.')
    for col in ('E', 'H', 'K'))


# ==========================================================================
# Demos con pycel
# ==========================================================================
def demo(ruta):
    exc = C.compilador(ruta)
    pruebas = []

    def ev(ref):
        return exc.evaluate(ref)

    r_ing2 = "'%s'!C%d" % (H_PYG, P['ingresos'])
    r_neto2 = "'%s'!C%d" % (H_PYG, P['neto'])
    r_fijos2 = "'%s'!C%d" % (H_PYG, P['tot_fijos'])
    r_tickets = "'%s'!B%d" % (H_SUP, S['tickets'])

    # 1. el total de fijos es la suma de sus componentes
    total = ev(r_fijos2)
    partes = sum(ev("'%s'!C%d" % (H_PYG, f))
                 for f in range(P['personal'], P['financieros'] + 1))
    pruebas.append(('El total de costes fijos suma sus catorce renglones',
                    abs(total - partes) < 0.01,
                    'total %.2f vs suma %.2f' % (total, partes)))

    # 2. el año 2 del P&L reproduce el año de crucero de datos_ejemplo
    esperado = D.cuenta_resultados_crucero()
    pruebas.append(('El año 2 del P&L es el año de crucero del juego de datos',
                    abs(ev(r_ing2) - esperado['ventas_sin_iva']) < 1.0,
                    'libro %.2f vs datos_ejemplo %.2f'
                    % (ev(r_ing2), esperado['ventas_sin_iva'])))

    # 3. el cuadro de amortización cierra
    pend = ev("'%s'!B%d" % (H_FIN, F_PENDIENTE))
    pruebas.append(('El cuadro de amortización devuelve todo el principal',
                    abs(pend) <= 0.01, 'pendiente al vencimiento %.2f' % pend))

    # 4. quitar el B2B mueve el punto muerto
    pm_con = ev("'%s'!B%d" % (H_CAN, K_PM_CON))
    pm_sin = ev("'%s'!B%d" % (H_CAN, K_PM_SIN))
    pruebas.append(('El punto muerto cambia al quitar el canal B2B',
                    abs(pm_con - pm_sin) > 1.0 and pm_sin < pm_con,
                    'con B2B %.2f, sin B2B %.2f' % (pm_con, pm_sin)))

    # 5. «sin dato» es cadena vacía, nunca cero
    exc.set_value(r_tickets, '')
    vacio = exc.evaluate(r_ing2)
    pruebas.append(('Sin tickets al día los ingresos salen VACÍOS, no a cero',
                    vacio == '', 'devuelve %r' % (vacio,)))

    # 6. los tres blancos declarados son cadena vacía, no un error
    blancos = [(h, c) for (h, c) in BLANCOS_OK]
    valores = [ev("'%s'!%s" % (h, c)) for h, c in blancos]
    pruebas.append(('Los pagos del 303 sin cuota van en BLANCO, no a cero',
                    all(v == '' for v in valores),
                    'devuelven %r' % (valores,)))

    # 7. la financiación nace cuadrada
    dif = ev("'%s'!B%d" % (H_FIN, F['diferencia']))
    pruebas.append(('El origen de fondos cubre la necesidad de caja de fábrica',
                    0 <= dif < 1.0, 'diferencia origen - usos %.2f €' % dif))

    # 8. el veredicto cambia al cambiar una entrada (compilador limpio: pycel
    #    no reevalúa bien un grafo al que ya se le ha inyectado un valor)
    otro = C.compilador(ruta)
    neto_base = otro.evaluate(r_neto2)
    otro.set_value(r_tickets, 120)
    neto_bajo = otro.evaluate(r_neto2)
    pruebas.append(('El resultado responde a la entrada de tickets al día',
                    neto_bajo < neto_base,
                    'con 120 tickets %.2f, con %d tickets %.2f'
                    % (neto_bajo, D.P('tickets_dia_crucero'), neto_base)))

    # 9. el año de crucero no se despega de la foto anual del juego de datos
    neto_libro = ev(r_neto2)
    rai_libro = ev("'%s'!C%d" % (H_PYG, P['rai']))
    esperado_rai = esperado['beneficio_neto']
    pruebas.append(('El resultado del año de crucero cuadra con el juego de '
                    'datos dentro del 2 %',
                    abs(rai_libro - esperado_rai) / abs(esperado_rai) < 0.02,
                    'libro %.2f vs datos_ejemplo %.2f (la diferencia son los '
                    'intereses: el libro le da al año 2 SUS intereses y '
                    'datos_ejemplo usa los del año 1); neto tras impuestos '
                    '%.2f' % (rai_libro, esperado_rai, neto_libro)))
    return pruebas


def main():
    D.gate_legal()
    wb = Workbook()
    wb.remove(wb.active)
    hoja_instrucciones(wb)
    hoja_supuestos(wb)
    hoja_inversion(wb)
    hoja_pyg(wb)
    hoja_equilibrio(wb)
    hoja_escenarios(wb)
    hoja_personal(wb)
    hoja_tesoreria(wb)
    hoja_financiacion(wb)
    hoja_canales(wb)
    res = C.cerrar(wb, NOMBRE, TITULO, mapa(), NOTAS_MAPA, BLANCOS_OK)
    C.resumen(NOMBRE, res, demo(res['ruta']))


if __name__ == '__main__':
    main()
