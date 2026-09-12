#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
gen_capacidad-obrador-y-local.py — libro 1 de «Cómo Montar una Pastelería»
(SPEC §2.2 fila 1; decisiones D17, D21, D22 y reglas de frontera R1, R5 y R6).

Hojas: Instrucciones · Parámetros · Zonas y m² · Capacidad por Equipo ·
Cuello de Botella · Ficha de Visita a Local.

QUÉ DECIDE ESTE LIBRO
---------------------
Si ESE local sirve, y cuántas piezas al día aguanta el equipo que estás
comprando. Son las dos preguntas que se contestan ANTES de firmar nada, y las
dos se contestan mal por el mismo motivo: se mira el precio del local y la
ficha del horno, y no se mira cuál de los dos manda.

LO QUE NO HACE, Y DÓNDE SE HACE (reglas de no-solape de la SPEC §2.3)
---------------------------------------------------------------------
* **R1 — No hay plan de producción semanal aquí.** Eso es
  `kit-tareas-pasteleria/10-plan-produccion-semanal.xlsx`, que organiza la
  producción en el obrador que YA tienes. Este libro decide la compra: cuántas
  piezas al día permite la máquina que estás a punto de pagar.
* **R5 — No repite el calendario de picos** (`BONUS-02-calendario-anual-tareas`
  del kit da las fechas y el qué). Aquí sólo entra el pico como MULTIPLICADOR
  de piezas al día, para ver si el equipo aguanta Reyes. Los euros del pico son
  del libro 3.
* **R6 — No dimensiona la plantilla.** Las horas de turno entran como
  parámetro; cuánta gente hace falta y lo que cuesta con la Seguridad Social
  dentro es el libro 8.
* El precio de cada equipo y su plazo de entrega son de los libros 2 y 7. Aquí
  el equipo entra sólo por su CAPACIDAD.

DECISIONES TÉCNICAS
-------------------
* **Un solo modelo de capacidad para los diez equipos**, y por eso se pueden
  comparar: `capacidad por ciclo × piezas por unidad de capacidad` = piezas por
  ciclo; `INT(minutos disponibles / minutos por ciclo)` = ciclos al día; y el
  **porcentaje del surtido que pasa por ese equipo** convierte su producción en
  las piezas de obrador que permite. Sin ese porcentaje, un abatidor por el que
  pasa el 40 % del surtido parecería limitar el obrador entero a lo que él
  abate, que es el error de bulto de cualquier cuenta de capacidad.
* El **cuello de botella es un `MIN` sobre una columna visible**, no un
  `SUMPRODUCT` que nadie puede auditar, y el equipo que limita sale con
  `INDEX/MATCH` sobre esa misma columna. Los equipos que no tienes devuelven
  `""` y `MIN` los ignora: por eso «¿lo tienes?» cambia el veredicto.
* **La ficha de visita distingue el dato medido del juicio.** Cuatro ítems
  tienen umbral numérico calculado desde «Parámetros» (caudal de campana,
  potencia, altura libre y superficie útil) y se resuelven solos; el resto son
  Sí/No/Por comprobar. Un solo «No» en un ítem eliminatorio manda «DESCARTAR».
* **La prueba de humos se cita por el RITE, no por el CTE DB-HS 3** (PA-07 y
  PA-07b de la verificación legal del 10-sep): el DB-HS 3 se aplica a viviendas
  y, en edificios de otro uso, sólo a aparcamientos y garajes; para un local
  remite al RITE. La regla de los 3 m es de vivienda. Escribir aquí «DB-HS 3»
  sería citar la norma equivocada en la casilla que más locales tumba.
* Cero constantes dentro de las fórmulas de cálculo; `IFERROR(...,"")` en toda
  división; «sin dato» = `""`; semáforos con `ISNUMBER`; desplegables contra
  RANGO; prohibidas `INDIRECT`, `COUNTA`, `PMT`, `OFFSET`, `XLOOKUP`, `LET`,
  `LAMBDA`, `RANK`, `NETWORKDAYS` e `IRR`: cero usos. Textos WinAnsi (cp1252).

Salida fija: build/capacidad-obrador-y-local.xlsx + mapa-capacidad-obrador-y-local.json
Via: Claude Code
"""
import json
import os
import subprocess
import sys

import openpyxl
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

AQUI = os.path.dirname(os.path.abspath(__file__))
RAIZ = os.path.normpath(os.path.join(AQUI, '..'))
sys.path.insert(0, os.path.join(RAIZ, 'guias-v2_0'))
sys.path.insert(0, AQUI)

import motor                                                   # noqa: E402
import datos_ejemplo as D                                      # noqa: E402
import _comun_pasteleria as C                                  # noqa: E402

motor.CTX['producto'] = C.PID

NOMBRE = 'capacidad-obrador-y-local'
TITULO = 'Capacidad del obrador y ficha de visita al local'

H_PAR = 'Parámetros'
H_ZON = 'Zonas y m2'
H_CAP = 'Capacidad por Equipo'
H_CUE = 'Cuello de Botella'
H_LOC = 'Ficha de Visita a Local'

#: Ids legales que este libro cita en celda. Se comprueban ANTES de escribir.
IDS_LEGALES = {
    'PA-03b': ('El obrador es zona inaccesible al público', 1),
    # PA-07 se retiró de aquí en la corrección del 2026-09-10 (hallazgo B2):
    # su `dato` dice que el CTE DB-HS 3 NO se aplica a un obrador, así que
    # citarlo en la celda de los humos leía como una verificación positiva
    # del CTE. La base correcta es PA-07b (RITE), abajo.
    'PA-07b': ('El aire de la campana es AE4 y no comparte conducto', 1),
    'PA-08': ('Horno de gas: inspección periódica cada cinco años', 1),
    'PA-12': ('Los 4 °C de la fila 9 del art. 4.1', 1),
    'PA-13': ('El abatidor del art. 5: un arcón no vale', 1),
    'PA-40': ('Accesibilidad: RD 193/2023', 1),
}

SI_NO = ('Sí', 'No')
RESPUESTAS = ('Sí', 'No', 'Por comprobar')
ENERGIAS = ('Eléctrica', 'Gas')

# --------------------------------------------------------------------------
# Parámetros del libro: (etiqueta, valor, unidad, procedencia, nota)
# --------------------------------------------------------------------------
P_INI = 6
PARAMS_LIBRO = [
    ('Horas de turno del obrador', D.NEGOCIO['horas_turno'], 'h por turno',
     D.NEGOCIO['fuente_horas_turno'],
     'Duración del turno de obrador. Es el número que multiplica todo lo demás: '
     'con dos turnos, la capacidad se duplica sin comprar una sola máquina.'),
    ('Turnos de obrador al día', 1, 'turnos', 'supuesto',
     'La Clara trabaja a un turno de madrugada. El segundo turno es la palanca '
     'más barata que existe para subir capacidad, y la más cara en personal.'),
    ('Horas productivas sobre la jornada', D.P('ratio_horas_productivas'), '%',
     'supuesto',
     'Parte de la jornada que la máquina está de verdad trabajando. El resto es '
     'recepción, limpieza, cambio de bandejas y esperas.'),
    ('Minutos disponibles al día', None, 'minutos', 'Se calcula',
     'Horas de turno × turnos × horas productivas × 60. De aquí salen los ciclos '
     'al día de cada equipo.'),
    ('Días de apertura a la semana', D.NEGOCIO['dias_apertura_semana'], 'días',
     D.NEGOCIO['fuente_dias_apertura'], D.NEGOCIO['nota_dias_apertura']),
    ('Días de apertura al año', D.NEGOCIO['dias_apertura_anio'], 'días',
     D.NEGOCIO['fuente_dias_apertura'],
     '312 días de apertura menos 12 de cierre por vacaciones.'),
    ('Piezas al día en velocidad de crucero', round(D.piezas_dia_crucero()),
     'piezas al día', 'supuesto',
     'Lo que tienes que sacar todos los días una vez rodado el negocio. Sale de '
     'los clientes al día por las piezas por ticket, los dos supuestos: no '
     'existe dato público de tráfico de pastelería. Si ya tienes el Kit de '
     'Tareas Pastelería, sustitúyelo por tu dato real de '
     'kit-tareas-pasteleria/10-plan-produccion-semanal.xlsx.'),
    ('Margen de seguridad exigido sobre la capacidad', 0.10, '%', 'supuesto',
     'Holgura mínima que quieres tener sobre el día normal. Por debajo de ella, '
     'el libro avisa: una avería, una baja o un pedido grande te dejan sin '
     'género aunque en teoría llegues.'),
    ('Superficie total del local', D.NEGOCIO['m2_total'], 'm2',
     D.NEGOCIO['fuente_m2'],
     'Los 90 m2 del caso de referencia. Es la superficie sobre la que casan los '
     'escenarios de obra de 600, 900 y 1.200 EUR/m2 del libro 2.'),
    ('Parte mínima del local dedicada a obrador', 0.45, '%', 'supuesto',
     'Criterio de la casa, no norma: por debajo de esto la tienda se come al '
     'obrador y acabas comprando semielaborados. Si vendes sobre todo por '
     'encargo, bájalo; si el despacho es tu escaparate, súbelo.'),
    ('Potencia eléctrica de los hornos', D.NEGOCIO['potencia_hornos_kw'], 'kW',
     D.NEGOCIO['fuente_potencia_hornos'],
     'PS-86a da 20-80 kW SOLO para los hornos. 34 kW es el punto de trabajo de '
     'los dos hornos de convección de La Clara.'),
    ('Potencia instalada total prevista', D.NEGOCIO['potencia_instalada_kw'],
     'kW', D.NEGOCIO['fuente_potencia_instalada'], D.NEGOCIO['nota_potencia']),
    ('Energía de los hornos', 'Eléctrica', 'Eléctrica o Gas', 'supuesto',
     'Cambia el caudal de campana que necesitas y, si eliges gas, activa la '
     'inspección periódica obligatoria de la instalación receptora.'),
    ('Caudal de extracción por kW, equipos eléctricos', 40, 'm3/h por kW',
     'PS-86b', 'Regla de dimensionado de campana para equipos eléctricos.'),
    ('Caudal de extracción por kW, equipos de gas', 50, 'm3/h por kW',
     'PS-86b', 'Regla de dimensionado de campana para equipos de gas.'),
    ('Caudal de campana necesario', None, 'm3/h', 'Se calcula',
     'Potencia de los hornos × el caudal por kW que corresponda a tu energía. '
     'Es el número que hay que llevar a la visita: si el local no puede darlo, '
     'no hay pastelería.'),
    ('Altura libre mínima que exiges al local', 3.0, 'm', 'supuesto',
     'Con campana, conducto y falso techo, por debajo de 3 m el obrador se '
     'queda sin aire y sin sitio para el conducto.'),
    ('Número de zonas del local', len(D.ZONAS), 'zonas', 'PS-86d',
     'Las zonas obligatorias de un obrador de pastelería. Si tu local no '
     'permite las siete, no es que quede apretado: es que falta una.'),
]
# Índices (0-based) de los parámetros que son fórmula.
P_MIN_DIA = 3
P_CAUDAL = 15
P_FILA = {}          # etiqueta -> fila


def fila_par(i):
    return P_INI + i


# --------------------------------------------------------------------------
# Equipos que marcan capacidad. Todo lo que no es capacidad (precio, plazo,
# IVA) vive en los libros 2 y 7: aquí el equipo entra sólo por lo que produce.
#
# (nombre, categoría, ¿lo tienes?, capacidad por ciclo, unidad, piezas por
#  unidad, minutos por ciclo, % del surtido, procedencia, nota)
# --------------------------------------------------------------------------
C_INI = 6
EQUIPOS_CAPACIDAD = [
    ('Amasadora de espiral, 25 kg', 'Amasado y batido', 'Sí',
     25, 'kg de masa por amasada', 19.6, 30, 0.55, 'supuesto',
     'Las piezas por kg salen de la tanda de croissant de la carta: 1,02 kg de '
     'masa para 20 piezas, es decir 19,6 piezas por kg. Cambia el número si tu '
     'surtido pesa más por pieza.'),
    ('Batidora planetaria, 20 L', 'Amasado y batido', 'Sí',
     20, 'L de masa o crema por batida', 8, 25, 0.30, 'PS-81',
     'Sammic BE-20, 900 W. Es la que marca el ritmo de cremas, mousses y '
     'merengues, no de las masas.'),
    ('Laminadora automática de masa, hasta 40 cm', 'Laminación', 'Sí',
     3, 'pastones por ciclo', 20, 20, 0.25, 'PS-79',
     'Sammic DF-40. Un pastón da una tanda de 20 piezas de bollería, así que '
     'tres pastones son 60 piezas por ciclo.'),
    ('Horno de convección, 6 bandejas', 'Calor', 'Sí',
     6, 'bandejas por hornada', 12, 25, 0.60, 'PS-72',
     'UNOX XB693 Bakerlux. Las piezas por bandeja son de bollería mediana: '
     'bájalas para tartas y súbelas para petit four.'),
    ('Horno de convección, 4 bandejas', 'Calor', 'Sí',
     4, 'bandejas por hornada', 12, 25, 0.25, 'PS-72',
     'SMEG ALFA420E1HDS. El segundo horno no está para producir más: está para '
     'poder cocer a dos temperaturas a la vez.'),
    ('Cámara de fermentación controlada (roll-in)', 'Frío positivo', 'Sí',
     2, 'carros por ciclo', 90, 120, 0.35, 'PS-83',
     'Es lo que permite que la bollería esté horneada a las 7:00 sin entrar a '
     'las 3. Su ciclo es largo, así que limita más de lo que parece.'),
    ('Abatidor de temperatura, 4 bandejas', 'Frío negativo', 'Sí',
     4, 'bandejas por ciclo', 12, 90, 0.40, 'PS-78',
     'Irinox, gama Multifresh. Ciclo largo y obligatorio para lo relleno: es el '
     'candidato número uno a cuello de botella de una pastelería. Un arcón '
     'doméstico NO vale (art. 5 del RD 1021/2022).'),
    ('Cámara frigorífica de conservación, 8 m3', 'Frío positivo', 'Sí',
     8, 'm3 útiles por rotación', 70, 400, 0.45, 'supuesto',
     'Aquí el "ciclo" es una rotación de la cámara: con 400 minutos entra una '
     'rotación al día. Es la que sostiene los 4 °C de lo relleno.'),
    ('Vitrina expositora pastelera refrigerada', 'Tienda y vitrina', 'Sí',
     0.93, 'm2 de exposición por reposición', 130, 120, 0.35, 'PS-76',
     'Docriluc VEPD-9-15-C, 0,93 m2 de exposición. Con 120 minutos por '
     'reposición salen tres reposiciones de vitrina al día.'),
    ('Horno de pisos modular (deck)', 'Calor', 'No',
     4, 'bandejas por hornada', 14, 35, 0.20, 'PS-73',
     'OPCIONAL. Marcado como "No" a propósito: cámbialo a "Sí" y mira cómo se '
     'mueve el cuello de botella antes de gastarte el dinero.'),
]

# --------------------------------------------------------------------------
# Ficha de visita: (ítem, eliminatorio, cómo se comprueba, respuesta por
#  defecto, dato medido o None, umbral (fórmula, None o texto), nota, id legal)
# --------------------------------------------------------------------------
L_INI = 6
VISITA = [
    ('El local puede tener salida de humos propia hasta cubierta', 'Sí',
     'Sube a la cubierta con el propietario y mira por dónde subiría el '
     'conducto. Pídelo por escrito antes de firmar.', 'Sí', None, None,
     'Referencia de sector, NO norma verificada (PS-86c): la salida suele '
     'tener que llegar a cubierta y rebasar la cumbrera en 1 m, pero la '
     'altura real la fija la ORDENANZA MUNICIPAL, no el CTE. Es la partida '
     'que más locales tumba y la que hay que resolver ANTES de firmar, no '
     'después.', None),
    ('La expulsión de la campana NO comparte conducto con el despacho ni con '
     'los aseos', 'Sí',
     'Pide el plano de instalaciones o abre el falso techo. Si sólo hay un '
     'patinillo, ya tienes la respuesta.', 'Sí', None, None,
     'El aire de la campana de humos es categoría AE4 del RITE: no puede '
     'recircularse, no puede usarse como aire de transferencia y su expulsión '
     'no puede ser común con la del aire del despacho o los aseos. Ojo: esto '
     'NO sale del CTE DB-HS 3, que se aplica a viviendas y, en otros usos, '
     'sólo a aparcamientos y garajes.', 'PA-07b'),
    ('El caudal de extracción disponible cubre el que necesitan tus hornos',
     'No',
     'Mide o pide la ficha del extractor existente. El umbral lo calcula la '
     'hoja «Parámetros» con tu potencia y tu tipo de energía.', 'Sí', 1400.0,
     'CAUDAL',
     'Regla de dimensionado: 40 m3/h por kW en equipos eléctricos y 50 en '
     'equipos de gas (PS-86b). Si el local se queda corto, el coste de '
     'sustituir el extractor va al libro 2.', None),
    ('La potencia eléctrica contratable llega a la que vas a instalar', 'Sí',
     'Pide el boletín eléctrico del local y consulta a la distribuidora la '
     'potencia disponible en la acometida.', 'Sí', 70.0, 'POTENCIA',
     'Subir la potencia disponible de una acometida puede costar más que la '
     'obra entera y depende de la distribuidora, no de tu presupuesto.', None),
    ('El planeamiento urbanístico admite la actividad en esa finca', 'Sí',
     'Consulta de compatibilidad urbanística en el ayuntamiento, por escrito y '
     'a nombre del local, no una llamada.', 'Sí', None, None,
     'Un local que hoy es un comercio no admite automáticamente un obrador: es '
     'actividad de transformación de alimentos, con humos y con horario de '
     'madrugada.', None),
    ('Los estatutos de la comunidad de propietarios no prohíben la actividad',
     'Sí',
     'Pide los estatutos inscritos y el acta de la última junta. Que el vecino '
     'de al lado sea un bar no basta.', 'Sí', None, None,
     'Una prohibición estatutaria inscrita se impone aunque el ayuntamiento te '
     'dé la licencia, y el conducto de humos suele pasar por zona común.',
     None),
    ('La altura libre del local llega a la que has fijado', 'No',
     'Metro laser desde el suelo terminado a la cara inferior del forjado, en '
     'el punto más bajo, no en el más alto.', 'Sí', 3.2, 'ALTURA',
     'Con campana, conducto y falso techo, cada centímetro cuenta. Mide donde '
     'va el horno.', None),
    ('La superficie útil llega a los m2 que has repartido por zonas', 'No',
     'Superficie útil real, sin contar muros ni patinillos. Compárala con el '
     'total de la hoja «Zonas y m2».', 'Sí', 92.0, 'SUPERFICIE',
     'La superficie de un anuncio suele ser construida. La que te sirve es la '
     'útil, y entre una y otra se van con facilidad 10 m2.', None),
    ('Se puede separar el obrador para que quede inaccesible al público', 'Sí',
     'Dibuja las siete zonas sobre el plano antes de firmar. Si no caben con '
     'la puerta donde está, no caben.', 'Sí', None, None,
     'El obrador es zona de elaboración inaccesible al público. No es una '
     'recomendación de diseño: es lo que separa un obrador de una cocina a la '
     'vista.', 'PA-03b'),
    ('Las siete zonas caben con marcha adelante y sin cruces', 'No',
     'Comprueba el recorrido en la hoja «Zonas y m2»: materia prima, obrador, '
     'producto terminado, venta.', 'Sí', None, None,
     'Que quepan los metros no quiere decir que quepa el recorrido. Un obrador '
     'donde la basura sale por donde entra la harina no pasa una inspección.',
     None),
    ('Hay acometida de agua fría y caliente y desagüe donde van los equipos',
     'No',
     'Localiza el desagüe existente. Mover un desagüe obliga a levantar el '
     'suelo del local entero.', 'Sí', None, None,
     'El horno con generador de vapor, el abatidor y el fregadero necesitan '
     'desagüe en su punto, no al fondo del local.', None),
    ('Suelos y paredes admiten acabado liso, lavable e impermeable', 'No',
     'Mira el estado real: azulejo roto, yeso, hormigón visto. El acabado del '
     'obrador es obra, no pintura.', 'Sí', None, None,
     'Superficies lisas, lavables e impermeables, encuentros pared-suelo '
     'redondeados y lavamanos no manual por zona (PS-86e). No es un extra: es '
     'requisito.', None),
    ('Caben aseo y vestuario de personal separados de los de clientes', 'No',
     'Cuenta los metros y mira dónde están los bajantes. Un aseo nuevo sin '
     'bajante cerca es obra grande.', 'Sí', None, None,
     'El personal necesita su aseo y su taquilla; el cliente, el suyo. Los 11 '
     'm2 de la hoja de zonas ya lo contemplan.', None),
    ('Se puede descargar género sin cruzar la zona de clientes', 'No',
     'Ve a la hora a la que descargarías, no a mediodía. Mira si hay carga y '
     'descarga y si el portal lo permite.', 'Sí', None, None,
     'La harina entra en sacos y a diario. Si entra por la puerta del '
     'despacho, entra por delante del cliente todas las mañanas.', None),
    ('La basura sale sin abrir al obrador', 'No',
     'Localiza el cuarto de residuos y su puerta. Debe estar cerrado y con '
     'contenedores identificados.', 'Sí', None, None,
     'El almacén de residuos no puede abrir al obrador. Es de los pocos '
     'defectos de plano que no tienen arreglo barato.', None),
    ('El local es accesible o admite ajustes razonables', 'No',
     'Mide el escalón de entrada y el ancho de la puerta. Un escalón en fachada '
     'puede depender del ayuntamiento, no de ti.', 'Sí', None, None,
     'El calendario del RD 193/2023 para titularidad privada es el 1-ene-2029 '
     'en bienes y servicios nuevos y el 1-ene-2030 en los existentes; el art. '
     '25.2 obliga sin umbral de superficie a garantizar la accesibilidad o a '
     'hacer ajustes razonables.', 'PA-40'),
    ('Si el horno es de gas, la instalación receptora está al día', 'Sí',
     'Pide el certificado de la última inspección periódica y la fecha de '
     'puesta en servicio.', 'Por comprobar', None, None,
     'Las instalaciones receptoras alimentadas desde red se inspeccionan cada '
     'cinco años, y el coste se repercute al usuario. Hasta 70 kW instalados '
     'la inspección incluye los aparatos de gas.', 'PA-08'),
    ('El contrato y las cargas del local están claros antes de firmar', 'Sí',
     'Nota simple del Registro, contrato de arrendamiento leído entero y quién '
     'paga la obra de adecuación.', 'Por comprobar', None, None,
     'Quién paga la obra, cuántos años dura el contrato y qué pasa con la obra '
     'al terminar son tres números del libro 2, no tres detalles.', None),
]


# --------------------------------------------------------------------------
def nl(idd):
    """Nota legal de la verificación del 10-sep, o aborta."""
    n = D.nota_legal(idd)
    if n is None:
        raise SystemExit('nota_legal(%s) devolvió None: gate_legal debería '
                         'haberlo cazado antes.' % idd)
    return n


def poner_nota(ws, coord, texto):
    from openpyxl.comments import Comment
    com = Comment(texto, 'AI Chef Pro')
    com.width = 420
    com.height = 130
    ws[coord].comment = com


NOTAS_LEGALES = {'n': 0}


def nota_legal_celda(ws, coord, idd):
    poner_nota(ws, coord, nl(idd))
    NOTAS_LEGALES['n'] += 1


# --------------------------------------------------------------------------
# Hoja «Instrucciones»
# --------------------------------------------------------------------------
PASOS = [
    '1. Hoja «Parámetros»: lo primero. Horas de turno, turnos al día, piezas '
    'que quieres sacar, superficie, potencia y tipo de energía de los hornos. '
    'Todo lo demás del libro se recalcula desde aquí; no hay ni un número '
    'escondido dentro de una fórmula.',
    '2. Hoja «Zonas y m2»: reparte los metros entre las siete zonas '
    'obligatorias y numera el recorrido de la marcha adelante, del 1 al 7. La '
    'hoja te dice si el reparto cuadra con la superficie declarada, qué parte '
    'del local es obrador y si el recorrido tiene un cruce.',
    '3. Hoja «Capacidad por Equipo»: una fila por máquina. Di si la tienes, '
    'cuánto cabe en un ciclo, cuántas piezas salen de esa cantidad, cuántos '
    'minutos dura el ciclo y qué parte de tu surtido pasa por ahí. La hoja '
    'devuelve las piezas al día que ESA máquina permite al obrador entero.',
    '4. Hoja «Cuello de Botella»: el resultado. Cuál es la máquina que manda, '
    'cuántas piezas al día aguanta el conjunto, cuánta holgura te queda sobre '
    'el día normal y qué pasa en cada uno de los seis picos del año. Aquí es '
    'donde se ve si Reyes te cabe en el obrador que estás comprando.',
    '5. Hoja «Ficha de Visita a Local»: imprímela y llévala a cada visita. '
    'Dieciocho comprobaciones, ocho de ellas eliminatorias. Un solo «No» en un '
    'ítem eliminatorio y el veredicto es DESCARTAR: no es pesimismo, es que ese '
    'defecto no se arregla con dinero razonable.',
]

NOTAS_LIBRO = [
    'ESTE LIBRO DECIDE UNA COMPRA, NO UNA SEMANA DE TRABAJO. La producción del '
    'día a día -qué se hace el martes y cuánto- vive en '
    'kit-tareas-pasteleria/10-plan-produccion-semanal.xlsx. Aquí sólo se '
    'contesta cuántas piezas al día permite el equipo que estás a punto de '
    'pagar, que es la pregunta que ya no se puede contestar después.',
    'EL PORCENTAJE DEL SURTIDO ES LA COLUMNA QUE MÁS SE OLVIDA. Un abatidor por '
    'el que pasa el 40 % de lo que haces y abate 192 piezas al día no limita tu '
    'obrador a 192 piezas: lo limita a 480. Sin esa columna, cualquier cuenta '
    'de capacidad sale mal y siempre por lo bajo.',
    'LOS PICOS NO SE RESUELVEN CON HORAS EXTRA. Reyes multiplica por seis la '
    'producción de un día normal. Ninguna plantilla absorbe eso: se absorbe con '
    'congelación programada, con producto estable a temperatura ambiente '
    'adelantado y con encargos cerrados con antelación. El calendario de los '
    'picos y el qué de cada uno está en el BONUS-02 del Kit de Tareas '
    'Pastelería; los euros de cada campaña, en el libro 3 de esta guía.',
    'LA PRUEBA DE HUMOS SE CITA POR EL RITE. El CTE DB-HS 3 se aplica a las '
    'viviendas y, en edificios de otro uso, sólo a aparcamientos y garajes: '
    'para un local remite al RITE, y la regla de los 3 m que circula por ahí es '
    'de vivienda. Lo que sí obliga en tu obrador es que el aire de la campana '
    '(categoría AE4) no comparta expulsión con el del despacho ni con el de los '
    'aseos.',
    'LOS PRECIOS Y LOS PLAZOS NO ESTÁN AQUÍ. Cuánto cuesta cada máquina y '
    'cuántas semanas tarda en llegar son los libros 2 (CAPEX) y 7 '
    '(equipamiento y proveedores). Este libro sólo mira lo que producen.',
]

CADENCIA = ('Cadencia: la hoja «Ficha de Visita a Local», UNA VEZ POR LOCAL que '
            'visites, y se guarda aunque descartes -la comparación entre tres '
            'locales es la que enseña qué estás aceptando-. Las hojas de '
            'capacidad, cada vez que cambies un equipo del presupuesto y, '
            'después de abrir, UNA VEZ AL AÑO antes de la campaña de Reyes.')


def hoja_instrucciones(wb):
    ws = wb.create_sheet('Instrucciones', 0)
    C.anchos(ws, {'A': 46, 'B': 46, 'C': 46})
    motor.val(ws, 'A1', TITULO)
    ws['A1'].font = Font(bold=True, size=16, color=C.ORO)
    ws.row_dimensions[1].height = 30
    motor.val(ws, 'A2', C.SUBTITULO)
    motor.val(ws, 'A3', 'Para qué sirve: decidir si ESE local sirve y cuántas '
                        'piezas al día aguanta el equipo que estás comprando.')
    ws['A3'].font = Font(italic=True, size=9)

    C.seccion(ws, 'A5', 'Instrucciones de uso')
    fila = 6
    for paso in PASOS:
        ws.merge_cells('A%d:C%d' % (fila, fila))
        motor.val(ws, 'A%d' % fila, paso, wrap=True)
        ws.row_dimensions[fila].height = 62
        fila += 1
    fila += 1
    motor.val(ws, 'A%d' % fila, C.NOTA_VERDES)
    ws['A%d' % fila].fill = PatternFill('solid', fgColor=motor.VERDE)
    fila += 2

    C.seccion(ws, 'A%d' % fila, 'Lo que conviene saber antes de empezar')
    fila += 1
    for texto in NOTAS_LIBRO:
        ws.merge_cells('A%d:C%d' % (fila, fila))
        motor.val(ws, 'A%d' % fila, texto, wrap=True)
        ws.row_dimensions[fila].height = 62
        fila += 1
    fila += 1

    C.seccion(ws, 'A%d' % fila, 'Dónde acaba este libro y dónde sigue el Kit '
                                'de Tareas Pastelería')
    fila += 1
    C.cabecera(ws, fila, [('A', 'Lo que necesitas'),
                          ('B', 'Dónde está'),
                          ('C', 'Por qué no está aquí')], altura=26)
    fila += 1
    FRONTERAS = [
        ('Qué se produce cada día de la semana y en qué orden',
         'kit-tareas-pasteleria/10-plan-produccion-semanal.xlsx',
         'Es la operación del obrador que YA tienes. Este libro decide la '
         'compra: cuántas piezas al día permite la máquina.'),
        ('Las fechas de cada campaña y qué se hace en cada una',
         'kit-tareas-pasteleria/BONUS-02-calendario-anual-tareas.xlsx',
         'El kit pone las fechas y el qué. Aquí el pico entra sólo como '
         'multiplicador de piezas al día; los euros de cada campaña son el '
         'libro 3 de esta guía.'),
        ('Cuánta gente hace falta y cuánto cuesta con la Seguridad Social',
         'Libro 8 de esta guía: plantilla-turnos-y-coste-personal.xlsx',
         'Aquí las horas de turno son un parámetro de capacidad, no un '
         'dimensionado de plantilla.'),
        ('Cuánto cuesta cada máquina y cuánto tarda en llegar',
         'Libros 2 y 7 de esta guía',
         'El precio y el plazo no cambian la capacidad. Se separan a propósito '
         'para poder decidir primero qué necesitas y después qué te cuesta.'),
    ]
    for que, donde, porque in FRONTERAS:
        motor.val(ws, 'A%d' % fila, que, wrap=True)
        motor.val(ws, 'B%d' % fila, donde, wrap=True)
        motor.val(ws, 'C%d' % fila, porque, wrap=True)
        ws.row_dimensions[fila].height = 54
        fila += 1
    ws.merge_cells('A%d:C%d' % (fila, fila))
    motor.val(ws, 'A%d' % fila,
              'Es COPIA declarada, no vínculo: en toda la guía no hay ni una '
              'fórmula que apunte a otro fichero. Un enlace entre libros se '
              'rompe en cuanto alguien mueve una carpeta, y entonces el libro '
              'miente sin avisar. Donde una celda verde traería un dato del '
              'Kit, viene con un valor por defecto propio, declarado como '
              'supuesto, y su nota te dice qué fichero del Kit lo sustituye.',
              wrap=True)
    ws['A%d' % fila].font = Font(italic=True, size=9)
    ws.row_dimensions[fila].height = 56
    fila += 2

    C.seccion(ws, 'A%d' % fila, 'Cada cuánto se usa este libro')
    fila += 1
    ws.merge_cells('A%d:C%d' % (fila, fila))
    motor.val(ws, 'A%d' % fila, CADENCIA, wrap=True)
    ws.row_dimensions[fila].height = 44
    fila += 2

    ws.merge_cells('A%d:C%d' % (fila, fila))
    motor.val(ws, 'A%d' % fila, C.NOTA_DESPROTEGER, wrap=True)
    fila += 2
    C.pie(ws, fila)
    C.pagina(ws, apaisado=False)
    return ws


# --------------------------------------------------------------------------
# Hoja «Parámetros»
# --------------------------------------------------------------------------
def hoja_parametros(wb):
    ws = wb.create_sheet(H_PAR)
    C.anchos(ws, {'A': 46, 'B': 14, 'C': 18, 'D': 18, 'E': 78})
    C.encabezar(ws, 'Parámetros', 'Todo lo que el libro calcula sale de aquí. '
                                  'Ni un número vive dentro de una fórmula.',
                col_fin='E')
    C.cabecera(ws, 5, [('A', 'Parámetro'), ('B', 'Valor'), ('C', 'Unidad'),
                       ('D', 'Procedencia'), ('E', 'Qué es y de dónde sale')])
    for i, (etq, valor, unidad, fuente, notatxt) in enumerate(PARAMS_LIBRO):
        fila = fila_par(i)
        P_FILA[etq] = fila
        motor.val(ws, 'A%d' % fila, etq, wrap=True)
        if valor is None:
            pass                                   # se rellena más abajo
        elif isinstance(valor, float) and unidad == '%':
            motor.val(ws, 'B%d' % fila, valor, fmt=C.FMT_PCT, verde_=True)
        elif isinstance(valor, str):
            motor.val(ws, 'B%d' % fila, valor, verde_=True)
        elif isinstance(valor, int):
            motor.val(ws, 'B%d' % fila, valor, fmt=C.FMT_ENT, verde_=True)
        else:
            motor.val(ws, 'B%d' % fila, valor, fmt=C.FMT_DEC1, verde_=True)
        motor.val(ws, 'C%d' % fila, unidad)
        motor.val(ws, 'D%d' % fila, fuente)
        motor.val(ws, 'E%d' % fila, notatxt, wrap=True)
        ws.row_dimensions[fila].height = 42

    f_min = fila_par(P_MIN_DIA)
    f_caudal = fila_par(P_CAUDAL)
    f_horas, f_turnos, f_prod = fila_par(0), fila_par(1), fila_par(2)
    f_pot, f_ener = fila_par(10), fila_par(12)
    f_q_ele, f_q_gas = fila_par(13), fila_par(14)
    motor.f(ws, 'B%d' % f_min,
            motor.iferror('B%d*B%d*B%d*60' % (f_horas, f_turnos, f_prod)),
            fmt=C.FMT_ENT)
    motor.f(ws, 'B%d' % f_caudal,
            motor.iferror('B%d*IF(B%d="Gas",B%d,B%d)'
                          % (f_pot, f_ener, f_q_gas, f_q_ele)),
            fmt=C.FMT_ENT)
    nota_legal_celda(ws, 'A%d' % f_caudal, 'PA-07b')

    fila = fila_par(len(PARAMS_LIBRO)) + 1
    refs, fila = C.bloque_listas(ws, fila, [('Energía de los hornos', ENERGIAS)],
                                 col='A')
    C.dv_rango(ws, ['B%d' % f_ener], refs['Energía de los hornos'],
               'Energía de los hornos', 'Elige «Eléctrica» o «Gas».')
    C.parrafo(ws, fila,
              'Los valores marcados «supuesto» son de la pastelería de ejemplo '
              '«La Clara» (90 m2, obrador de 45 m2, un turno de madrugada). No '
              'son datos de sector: son un punto de partida coherente para que '
              'puedas ver el libro funcionando antes de meter los tuyos.',
              col_ini='A', col_fin='E', alto=42)
    C.pagina(ws, apaisado=True, titulos='5:5')
    return ws


# --------------------------------------------------------------------------
# Hoja «Zonas y m2»
# --------------------------------------------------------------------------
#: Paso del recorrido de marcha adelante, por nombre de zona (supuesto, verde).
ORDEN_MARCHA = {
    'Almacén de materia prima': 1,
    'Cámara frigorífica': 2,
    'Obrador': 3,
    'Almacén de producto terminado': 4,
    'Zona de venta (despacho)': 5,
    'Almacén de residuos': 6,
    'Vestuarios y aseos de personal': 7,
}
Z_FILA = {}


def hoja_zonas(wb):
    ws = wb.create_sheet(H_ZON)
    C.anchos(ws, {'A': 5, 'B': 34, 'C': 22, 'D': 11, 'E': 11, 'F': 11,
                  'G': 86})
    C.encabezar(ws, 'Zonas y metros cuadrados',
                'Las siete zonas obligatorias de un obrador de pastelería '
                '(PS-86d) y el recorrido de marcha adelante.', col_fin='G')
    C.cabecera(ws, 5, [('A', 'Nº'), ('B', 'Zona'), ('C', 'Bloque'),
                       ('D', 'm2'), ('E', '% del local'),
                       ('F', 'Paso del recorrido'),
                       ('G', 'Qué tiene que cumplir')])
    fila = 6
    for i, (zona, bloque, m2, _fuente, notatxt) in enumerate(D.ZONAS):
        Z_FILA[zona] = fila
        motor.val(ws, 'A%d' % fila, i + 1, fmt=C.FMT_ENT)
        motor.val(ws, 'B%d' % fila, zona, wrap=True)
        motor.val(ws, 'C%d' % fila, bloque)
        motor.val(ws, 'D%d' % fila, m2, fmt=C.FMT_DEC1, verde_=True)
        motor.f(ws, 'E%d' % fila, motor.iferror('D%d/$D$14' % fila),
                fmt=C.FMT_PCT)
        motor.val(ws, 'F%d' % fila, ORDEN_MARCHA[zona], fmt=C.FMT_ENT,
                  verde_=True)
        motor.val(ws, 'G%d' % fila, notatxt, wrap=True)
        ws.row_dimensions[fila].height = 44
        fila += 1
    z_ini, z_fin = 6, fila - 1
    if 'Cámara frigorífica' in Z_FILA:
        nota_legal_celda(ws, 'G%d' % Z_FILA['Cámara frigorífica'], 'PA-12')
    nota_legal_celda(ws, 'G%d' % Z_FILA['Obrador'], 'PA-03b')

    f_tot = 14
    motor.val(ws, 'B%d' % f_tot, 'TOTAL repartido', bold=True)
    motor.f(ws, 'D%d' % f_tot, '=SUM(D%d:D%d)' % (z_ini, z_fin),
            fmt=C.FMT_DEC1, bold=True)
    motor.f(ws, 'E%d' % f_tot, '=SUM(E%d:E%d)' % (z_ini, z_fin), fmt=C.FMT_PCT,
            bold=True)
    C.destacado(ws, 'D%d' % f_tot)

    f_dec = 15
    motor.val(ws, 'B%d' % f_dec, 'Superficie declarada en «Parámetros»')
    motor.f(ws, 'D%d' % f_dec, "='%s'!B%s" % (H_PAR, P_FILA[
        'Superficie total del local']), fmt=C.FMT_DEC1)
    f_desc = 16
    motor.val(ws, 'B%d' % f_desc, 'Descuadre (m2)')
    motor.f(ws, 'D%d' % f_desc, motor.iferror('D%d-D%d' % (f_tot, f_dec)),
            fmt=C.FMT_DEC1)
    motor.val(ws, 'G%d' % f_desc,
              'Si no es cero, o te sobran metros sin asignar o has repartido '
              'más de los que tiene el local.', wrap=True)
    motor.regla_expresion(ws, 'D%d' % f_desc,
                          '=AND(ISNUMBER($D$%d),ABS($D$%d)>0.05)'
                          % (f_desc, f_desc))

    f_obr, f_pobr = 18, 19
    motor.val(ws, 'B%d' % f_obr, 'Metros de obrador')
    motor.f(ws, 'D%d' % f_obr,
            '=SUMIF(C%d:C%d,"Obrador",D%d:D%d)' % (z_ini, z_fin, z_ini, z_fin),
            fmt=C.FMT_DEC1)
    motor.val(ws, 'B%d' % f_pobr, 'Parte del local dedicada a obrador')
    motor.f(ws, 'D%d' % f_pobr, motor.iferror('D%d/D%d' % (f_obr, f_tot)),
            fmt=C.FMT_PCT)
    f_ven, f_pven = 20, 21
    motor.val(ws, 'B%d' % f_ven, 'Metros de zona de venta (despacho)')
    motor.f(ws, 'D%d' % f_ven,
            '=SUMIF(C%d:C%d,"Despacho",D%d:D%d)' % (z_ini, z_fin, z_ini, z_fin),
            fmt=C.FMT_DEC1)
    motor.val(ws, 'B%d' % f_pven, 'Parte del local dedicada a venta')
    motor.f(ws, 'D%d' % f_pven, motor.iferror('D%d/D%d' % (f_ven, f_tot)),
            fmt=C.FMT_PCT)

    f_min_obr = 22
    motor.val(ws, 'B%d' % f_min_obr, 'Mínimo de obrador que te has exigido')
    motor.f(ws, 'D%d' % f_min_obr, "='%s'!B%s" % (H_PAR, P_FILA[
        'Parte mínima del local dedicada a obrador']), fmt=C.FMT_PCT)
    f_ver_rep = 23
    motor.val(ws, 'B%d' % f_ver_rep, 'Veredicto del reparto', bold=True)
    motor.f(ws, 'D%d' % f_ver_rep,
            '=IF(D%d="","",IF(D%d<D%d,"El obrador se queda corto",'
            '"Reparto suficiente"))' % (f_pobr, f_pobr, f_min_obr), bold=True)
    C.destacado(ws, 'D%d' % f_ver_rep)
    motor.semaforo_texto(ws, 'D%d' % f_ver_rep,
                         (('El obrador se queda corto', motor.CF_ROJO_BG,
                           motor.CF_ROJO_FG),
                          ('Reparto suficiente', motor.CF_VERDE_BG,
                           motor.CF_VERDE_FG)))

    C.seccion(ws, 'B25', 'Marcha adelante: la materia prima entra por un '
                         'extremo y el producto sale por el otro')
    f_sum, f_deb, f_num = 26, 27, 28
    motor.val(ws, 'B%d' % f_sum, 'Suma de los pasos que has numerado')
    motor.f(ws, 'D%d' % f_sum, '=SUM(F%d:F%d)' % (z_ini, z_fin), fmt=C.FMT_ENT)
    motor.val(ws, 'B%d' % f_deb, 'Suma que debería dar (1 + 2 + ... + n)')
    motor.f(ws, 'D%d' % f_deb,
            motor.iferror("'%s'!B%s*('%s'!B%s+1)/2"
                          % (H_PAR, P_FILA['Número de zonas del local'],
                             H_PAR, P_FILA['Número de zonas del local'])),
            fmt=C.FMT_ENT)
    motor.val(ws, 'B%d' % f_num, '¿Están numerados del 1 al n sin repetir?')
    motor.f(ws, 'D%d' % f_num,
            '=IF(AND(D%d=D%d,MIN(F%d:F%d)=1,MAX(F%d:F%d)=\'%s\'!B%s),"Sí",'
            '"Revisa la numeración")'
            % (f_sum, f_deb, z_ini, z_fin, z_ini, z_fin, H_PAR,
               P_FILA['Número de zonas del local']))
    motor.val(ws, 'G%d' % f_num,
              'Comprobación de numeración: la suma tiene que cuadrar y los '
              'pasos ir del 1 al número de zonas. No detecta una permutación '
              'que sume igual; lo que detecta -y es lo que pasa siempre- es '
              'un número repetido o uno saltado.', wrap=True)
    ws.row_dimensions[f_num].height = 44

    comprobaciones = [
        (29, 'La materia prima entra antes que el obrador',
         'F%d<F%d' % (Z_FILA['Almacén de materia prima'], Z_FILA['Obrador'])),
        (30, 'La cámara de conservación va antes que el obrador',
         'F%d<F%d' % (Z_FILA['Cámara frigorífica'], Z_FILA['Obrador'])),
        (31, 'El producto terminado sale después del obrador',
         'F%d>F%d' % (Z_FILA['Almacén de producto terminado'],
                      Z_FILA['Obrador'])),
        (32, 'El producto terminado llega a la venta sin volver atrás',
         'F%d<=F%d' % (Z_FILA['Almacén de producto terminado'],
                       Z_FILA['Zona de venta (despacho)'])),
        (33, 'Los residuos salen después del obrador',
         'F%d>F%d' % (Z_FILA['Almacén de residuos'], Z_FILA['Obrador'])),
    ]
    for fila_c, etq, cond in comprobaciones:
        motor.val(ws, 'B%d' % fila_c, etq, wrap=True)
        motor.f(ws, 'D%d' % fila_c, '=IF(%s,"Sí","No")' % cond)
        motor.semaforo_texto(ws, 'D%d' % fila_c,
                             (('No', motor.CF_ROJO_BG, motor.CF_ROJO_FG),
                              ('Sí', motor.CF_VERDE_BG, motor.CF_VERDE_FG)))
    f_ver_marcha = 35
    motor.val(ws, 'B%d' % f_ver_marcha, 'VEREDICTO DE MARCHA ADELANTE',
              bold=True)
    motor.f(ws, 'D%d' % f_ver_marcha,
            '=IF(COUNTIF(D%d:D%d,"No")>0,"Hay un cruce en el recorrido",'
            '"Marcha adelante correcta")' % (29, 33), bold=True)
    C.destacado(ws, 'D%d' % f_ver_marcha)
    motor.semaforo_texto(ws, 'D%d' % f_ver_marcha,
                         (('Hay un cruce en el recorrido', motor.CF_ROJO_BG,
                           motor.CF_ROJO_FG),
                          ('Marcha adelante correcta', motor.CF_VERDE_BG,
                           motor.CF_VERDE_FG)))
    C.parrafo(ws, 37,
              'La marcha adelante no es una preferencia de diseño: es que la '
              'materia prima sucia y el producto terminado no se crucen nunca. '
              'Cuando el local no la permite, no se arregla con un cartel ni '
              'con un horario; se arregla con obra o no se arregla. Por eso '
              'este recorrido se dibuja sobre el plano ANTES de firmar.',
              col_ini='B', col_fin='G', alto=44)
    C.pagina(ws, apaisado=True, titulos='5:5')
    return ws


# --------------------------------------------------------------------------
# Hoja «Capacidad por Equipo»
# --------------------------------------------------------------------------
def hoja_capacidad(wb):
    ws = wb.create_sheet(H_CAP)
    C.anchos(ws, {'A': 5, 'B': 40, 'C': 18, 'D': 12, 'E': 12, 'F': 26,
                  'G': 12, 'H': 12, 'I': 12, 'J': 11, 'K': 13, 'L': 13,
                  'M': 15, 'N': 15, 'O': 14, 'P': 74})
    C.encabezar(ws, 'Capacidad por equipo',
                'Un solo modelo para todas las máquinas: lo que cabe en un '
                'ciclo, cuántas piezas salen de eso, cuánto dura el ciclo y '
                'qué parte de tu surtido pasa por ahí.', col_fin='P')
    C.cabecera(ws, 5, [
        ('A', 'Nº'), ('B', 'Equipo'), ('C', 'Categoría'),
        ('D', '¿Lo tienes?'), ('E', 'Capacidad por ciclo'),
        ('F', 'Unidad de la capacidad'), ('G', 'Piezas por unidad'),
        ('H', 'Piezas por ciclo'), ('I', 'Minutos por ciclo'),
        ('J', 'Ciclos al día'), ('K', 'Piezas al día del equipo'),
        ('L', '% del surtido que pasa'),
        ('M', 'Piezas/día de obrador que permite'),
        ('N', 'Cuenta para el cuello'), ('O', 'Procedencia'), ('P', 'Nota')],
        altura=54)
    fila = C_INI
    f_min_dia = fila_par(P_MIN_DIA)
    for i, (nombre, cat, tiene, cap, unidad, ppu, minutos, pct, fuente,
            notatxt) in enumerate(EQUIPOS_CAPACIDAD):
        motor.val(ws, 'A%d' % fila, i + 1, fmt=C.FMT_ENT)
        motor.val(ws, 'B%d' % fila, nombre, wrap=True)
        motor.val(ws, 'C%d' % fila, cat)
        motor.val(ws, 'D%d' % fila, tiene, verde_=True)
        motor.val(ws, 'E%d' % fila, cap, fmt=C.FMT_DEC1, verde_=True)
        motor.val(ws, 'F%d' % fila, unidad, wrap=True)
        motor.val(ws, 'G%d' % fila, ppu, fmt=C.FMT_DEC1, verde_=True)
        motor.f(ws, 'H%d' % fila, motor.iferror('E%d*G%d' % (fila, fila)),
                fmt=C.FMT_ENT)
        motor.val(ws, 'I%d' % fila, minutos, fmt=C.FMT_ENT, verde_=True)
        motor.f(ws, 'J%d' % fila,
                motor.iferror("INT('%s'!$B$%d/I%d)" % (H_PAR, f_min_dia, fila)),
                fmt=C.FMT_ENT)
        motor.f(ws, 'K%d' % fila, motor.iferror('H%d*J%d' % (fila, fila)),
                fmt=C.FMT_ENT)
        motor.val(ws, 'L%d' % fila, pct, fmt=C.FMT_PCT, verde_=True)
        motor.f(ws, 'M%d' % fila, motor.iferror('K%d/L%d' % (fila, fila)),
                fmt=C.FMT_ENT)
        motor.f(ws, 'N%d' % fila,
                '=IF(D%d="Sí",M%d,"")' % (fila, fila), fmt=C.FMT_ENT)
        motor.val(ws, 'O%d' % fila, fuente)
        motor.val(ws, 'P%d' % fila, notatxt, wrap=True)
        ws.row_dimensions[fila].height = 46
        fila += 1
    c_fin = fila - 1
    nota_legal_celda(ws, 'B%d' % (C_INI + 6), 'PA-13')
    nota_legal_celda(ws, 'B%d' % (C_INI + 7), 'PA-12')

    motor.semaforo_isnumber(ws, 'J%d:J%d' % (C_INI, c_fin),
                            '$J%d' % C_INI, operador='<', umbral='1')

    f = c_fin + 2
    motor.val(ws, 'B%d' % f, 'Minutos disponibles al día (de «Parámetros»)')
    motor.f(ws, 'E%d' % f, "='%s'!B%d" % (H_PAR, f_min_dia), fmt=C.FMT_ENT)
    motor.val(ws, 'B%d' % (f + 1), 'Equipos que dices tener')
    motor.f(ws, 'E%d' % (f + 1), '=COUNTIF(D%d:D%d,"Sí")' % (C_INI, c_fin),
            fmt=C.FMT_ENT)
    motor.val(ws, 'B%d' % (f + 2), 'Equipos que no cuentan (ciclo más largo '
                                   'que la jornada)')
    motor.f(ws, 'E%d' % (f + 2),
            '=COUNTIF(J%d:J%d,0)' % (C_INI, c_fin), fmt=C.FMT_ENT)
    motor.val(ws, 'P%d' % (f + 2),
              'Un cero en «ciclos al día» quiere decir que el ciclo de esa '
              'máquina no cabe en la jornada que has declarado. No es un '
              'error del libro: es que con ese turno esa máquina no produce '
              'nada.', wrap=True)
    ws.row_dimensions[f + 2].height = 44

    fila = f + 4
    refs, fila = C.bloque_listas(ws, fila, [('¿Lo tienes?', SI_NO)], col='B')
    C.dv_rango(ws, ['D%d' % r for r in range(C_INI, c_fin + 1)],
               refs['¿Lo tienes?'], '¿Lo tienes?',
               'Elige «Sí» o «No» de la lista.')
    C.parrafo(ws, fila,
              'Las columnas verdes de esta hoja son las que tienes que pelear '
              'con el comercial: cuánto cabe de verdad en un ciclo, cuánto '
              'dura el ciclo completo (con carga y descarga, no el que pone el '
              'folleto) y qué parte de tu surtido pasa por ahí. Si ya tienes el '
              'Kit de Tareas Pastelería, el porcentaje real de cada máquina lo '
              'sacas de kit-tareas-pasteleria/10-plan-produccion-semanal.xlsx; '
              'mientras tanto, los valores por defecto son supuestos coherentes '
              'con la carta de 30 referencias de «La Clara».',
              col_ini='B', col_fin='P', alto=56)
    C.pagina(ws, apaisado=True, titulos='5:5')
    return ws


# --------------------------------------------------------------------------
# Hoja «Cuello de Botella»
# --------------------------------------------------------------------------
K_OBJ, K_CAP, K_EQU, K_HOL, K_HOLP, K_MAR, K_VER = 6, 7, 8, 9, 10, 11, 12
K_SEG, K_SEGE = 13, 14
K_PICO_CAB = 17
K_PICO_INI = 18


def hoja_cuello(wb):
    ws = wb.create_sheet(H_CUE)
    C.anchos(ws, {'A': 44, 'B': 20, 'C': 16, 'D': 16, 'E': 16, 'F': 15,
                  'G': 14, 'H': 26, 'I': 15, 'J': 74})
    C.encabezar(ws, 'Cuello de botella',
                'El equipo que manda, las piezas al día que aguanta el '
                'conjunto y qué pasa en cada uno de los seis picos del año.',
                col_fin='J')
    n_ini, n_fin = C_INI, C_INI + len(EQUIPOS_CAPACIDAD) - 1
    rng_n = "'%s'!$N$%d:$N$%d" % (H_CAP, n_ini, n_fin)
    rng_b = "'%s'!$B$%d:$B$%d" % (H_CAP, n_ini, n_fin)

    C.seccion(ws, 'A5', 'El día normal')
    motor.val(ws, 'A%d' % K_OBJ, 'Piezas al día que quieres sacar')
    motor.f(ws, 'B%d' % K_OBJ, "='%s'!B%d"
            % (H_PAR, fila_par(6)), fmt=C.FMT_ENT)
    motor.val(ws, 'A%d' % K_CAP, 'Piezas al día que permite el conjunto')
    motor.f(ws, 'B%d' % K_CAP, motor.iferror('MIN(%s)' % rng_n), fmt=C.FMT_ENT,
            bold=True)
    C.destacado(ws, 'B%d' % K_CAP)
    motor.val(ws, 'A%d' % K_EQU, 'El equipo que limita', bold=True)
    motor.f(ws, 'B%d' % K_EQU,
            motor.iferror('INDEX(%s,MATCH(B%d,%s,0))' % (rng_b, K_CAP, rng_n)),
            bold=True)
    C.destacado(ws, 'B%d' % K_EQU)
    motor.val(ws, 'A%d' % K_HOL, 'Holgura sobre el día normal (piezas)')
    motor.f(ws, 'B%d' % K_HOL, motor.iferror('B%d-B%d' % (K_CAP, K_OBJ)),
            fmt=C.FMT_ENT)
    motor.val(ws, 'A%d' % K_HOLP, 'Holgura sobre el día normal (%)')
    motor.f(ws, 'B%d' % K_HOLP, motor.iferror('B%d/B%d' % (K_HOL, K_OBJ)),
            fmt=C.FMT_PCT)
    motor.val(ws, 'A%d' % K_MAR, 'Margen de seguridad que te has exigido')
    motor.f(ws, 'B%d' % K_MAR, "='%s'!B%d" % (H_PAR, fila_par(7)),
            fmt=C.FMT_PCT)
    motor.val(ws, 'A%d' % K_VER, 'VEREDICTO DEL DÍA NORMAL', bold=True)
    motor.f(ws, 'B%d' % K_VER,
            '=IF(B%d="","",IF(B%d<B%d,"NO LLEGAS",IF(B%d<B%d,"AJUSTADO",'
            '"SUFICIENTE")))' % (K_CAP, K_CAP, K_OBJ, K_HOLP, K_MAR),
            bold=True)
    C.destacado(ws, 'B%d' % K_VER)
    motor.semaforo_texto(ws, 'B%d' % K_VER,
                         (('NO LLEGAS', motor.CF_ROJO_BG, motor.CF_ROJO_FG),
                          ('AJUSTADO', motor.CF_AMBAR_BG, motor.CF_AMBAR_FG),
                          ('SUFICIENTE', motor.CF_VERDE_BG,
                           motor.CF_VERDE_FG)))
    motor.val(ws, 'J%d' % K_VER,
              '«AJUSTADO» quiere decir que pasas el día normal, pero por '
              'debajo del margen de seguridad que tú mismo has pedido: una '
              'avería, una baja o un encargo grande te dejan sin género.',
              wrap=True)
    ws.row_dimensions[K_VER].height = 44

    motor.val(ws, 'A%d' % K_SEG, 'El segundo equipo más justo permite')
    motor.f(ws, 'B%d' % K_SEG, motor.iferror('SMALL(%s,2)' % rng_n),
            fmt=C.FMT_ENT)
    motor.val(ws, 'A%d' % K_SEGE, 'Y es')
    motor.f(ws, 'B%d' % K_SEGE,
            motor.iferror('INDEX(%s,MATCH(B%d,%s,0))' % (rng_b, K_SEG, rng_n)))
    motor.val(ws, 'J%d' % K_SEGE,
              'Mira los dos juntos antes de comprar: si el segundo está muy '
              'cerca del primero, resolver sólo el cuello de botella te sube '
              'la capacidad cuatro piezas y te deja el dinero fuera.',
              wrap=True)
    ws.row_dimensions[K_SEGE].height = 44

    C.seccion(ws, 'A16', 'Los seis picos del año')
    C.cabecera(ws, K_PICO_CAB, [
        ('A', 'Pico'), ('B', 'Fechas'), ('C', 'Factor sobre el día normal'),
        ('D', 'Piezas/día que hacen falta'),
        ('E', 'Piezas/día que permite el conjunto'),
        ('F', 'Déficit (piezas/día)'), ('G', '¿Aguantas?'),
        ('H', 'Producto estrella'),
        ('I', 'Uds/día del estrella en campaña'), ('J', 'Nota')], altura=54)
    fila = K_PICO_INI
    for p in D.PICOS:
        motor.val(ws, 'A%d' % fila, p['nombre'])
        motor.val(ws, 'B%d' % fila, p['fechas'])
        motor.val(ws, 'C%d' % fila, p['factor_obrador'], fmt=C.FMT_DEC1,
                  verde_=True)
        motor.f(ws, 'D%d' % fila, motor.iferror('$B$%d*C%d' % (K_OBJ, fila)),
                fmt=C.FMT_ENT)
        motor.f(ws, 'E%d' % fila, motor.iferror('$B$%d' % K_CAP),
                fmt=C.FMT_ENT)
        motor.f(ws, 'F%d' % fila, motor.iferror('E%d-D%d' % (fila, fila)),
                fmt=C.FMT_ENT)
        motor.f(ws, 'G%d' % fila,
                '=IF(F%d="","",IF(F%d<0,"NO","SÍ"))' % (fila, fila))
        motor.val(ws, 'H%d' % fila, p['producto_estrella'])
        motor.val(ws, 'I%d' % fila, p['uds_dia_pico'], fmt=C.FMT_ENT,
                  verde_=True)
        motor.val(ws, 'J%d' % fila, p['nota'], wrap=True)
        ws.row_dimensions[fila].height = 62
        fila += 1
    p_fin = fila - 1
    motor.semaforo_isnumber(ws, 'F%d:F%d' % (K_PICO_INI, p_fin),
                            '$F%d' % K_PICO_INI, operador='<', umbral='0')
    motor.semaforo_texto(ws, 'G%d:G%d' % (K_PICO_INI, p_fin),
                         (('NO', motor.CF_ROJO_BG, motor.CF_ROJO_FG),
                          ('SÍ', motor.CF_VERDE_BG, motor.CF_VERDE_FG)))

    f = p_fin + 2
    motor.val(ws, 'A%d' % f, 'Picos que NO aguantas')
    motor.f(ws, 'B%d' % f, '=COUNTIF(G%d:G%d,"NO")' % (K_PICO_INI, p_fin),
            fmt=C.FMT_ENT)
    motor.val(ws, 'A%d' % (f + 1), 'El peor déficit (piezas/día)')
    motor.f(ws, 'B%d' % (f + 1),
            motor.iferror('MIN(F%d:F%d)' % (K_PICO_INI, p_fin)), fmt=C.FMT_ENT)
    motor.val(ws, 'A%d' % (f + 2), 'Y es el pico de')
    motor.f(ws, 'B%d' % (f + 2),
            motor.iferror('INDEX($A$%d:$A$%d,MATCH(B%d,$F$%d:$F$%d,0))'
                          % (K_PICO_INI, p_fin, f + 1, K_PICO_INI, p_fin)))
    motor.val(ws, 'A%d' % (f + 3), 'VEREDICTO DE CAMPAÑA', bold=True)
    motor.f(ws, 'B%d' % (f + 3),
            '=IF(B%d=0,"Aguantas los seis picos",IF(B%d=1,'
            '"Un pico se te va del obrador","Varios picos se te van del '
            'obrador"))' % (f, f), bold=True)
    C.destacado(ws, 'B%d' % (f + 3))
    motor.semaforo_texto(ws, 'B%d' % (f + 3),
                         (('Aguantas los seis picos', motor.CF_VERDE_BG,
                           motor.CF_VERDE_FG),
                          ('Un pico se te va del obrador', motor.CF_AMBAR_BG,
                           motor.CF_AMBAR_FG),
                          ('Varios picos se te van del obrador',
                           motor.CF_ROJO_BG, motor.CF_ROJO_FG)))

    C.parrafo(ws, f + 5,
              'Que un pico no quepa en el obrador NO es motivo para comprar más '
              'máquina: comprar para el día de Reyes es tener el dinero parado '
              'los otros 299 días. Lo que se hace es adelantar producción '
              '(congelación programada del art. 5, con abatidor, no con un '
              'arcón), cargar la campaña en lo estable a temperatura ambiente, '
              'que se puede hacer con días de antelación, y cerrar encargos con '
              'fecha. Las fechas de cada campaña y su lista de tareas están en '
              'el BONUS-02 del Kit de Tareas Pastelería; los euros que mueve '
              'cada una, en el libro 3 de esta guía.',
              col_ini='A', col_fin='J', alto=62)
    C.parrafo(ws, f + 7,
              'La columna «uds/día del estrella en campaña» está aquí como '
              'contraste, no como cálculo: son las unidades del producto que '
              'define el pico. Los 420 roscones al día de Reyes frente a los '
              'cero de un día normal son la razón por la que esta hoja existe.',
              col_ini='A', col_fin='J', alto=44)
    C.pagina(ws, apaisado=True, titulos='17:17')
    return ws


# --------------------------------------------------------------------------
# Hoja «Ficha de Visita a Local»
# --------------------------------------------------------------------------
def hoja_local(wb):
    ws = wb.create_sheet(H_LOC)
    C.anchos(ws, {'A': 5, 'B': 52, 'C': 13, 'D': 46, 'E': 15, 'F': 13,
                  'G': 13, 'H': 13, 'I': 15, 'J': 80})
    C.encabezar(ws, 'Ficha de visita a local',
                'Imprímela y llévala a cada visita. Un solo «No» en un ítem '
                'eliminatorio y el veredicto es DESCARTAR.', col_fin='J')
    C.cabecera(ws, 5, [
        ('A', 'Nº'), ('B', 'Qué se comprueba'), ('C', '¿Eliminatorio?'),
        ('D', 'Cómo se comprueba'), ('E', 'Respuesta'), ('F', 'Dato medido'),
        ('G', 'Umbral'), ('H', '¿Cumple el dato?'), ('I', 'Veredicto'),
        ('J', 'Nota y norma')], altura=44)

    umbrales = {
        'CAUDAL': ("='%s'!B%d" % (H_PAR, fila_par(P_CAUDAL)), C.FMT_ENT),
        'POTENCIA': ("='%s'!B%d" % (H_PAR, fila_par(11)), C.FMT_DEC1),
        'ALTURA': ("='%s'!B%d" % (H_PAR, fila_par(16)), C.FMT_DEC1),
        'SUPERFICIE': ("='%s'!B%d" % (H_PAR, fila_par(8)), C.FMT_DEC1),
    }
    fila = L_INI
    for i, (item, elim, como, resp, medido, umbral, notatxt,
            idlegal) in enumerate(VISITA):
        motor.val(ws, 'A%d' % fila, i + 1, fmt=C.FMT_ENT)
        motor.val(ws, 'B%d' % fila, item, wrap=True)
        motor.val(ws, 'C%d' % fila, elim)
        motor.val(ws, 'D%d' % fila, como, wrap=True)
        motor.val(ws, 'E%d' % fila, resp, verde_=True)
        if medido is not None:
            motor.val(ws, 'F%d' % fila, medido, fmt=umbrales[umbral][1],
                      verde_=True)
            motor.f(ws, 'G%d' % fila, umbrales[umbral][0],
                    fmt=umbrales[umbral][1])
            motor.f(ws, 'H%d' % fila,
                    '=IF(AND(ISNUMBER(F%d),ISNUMBER(G%d)),IF(F%d>=G%d,"Sí",'
                    '"No"),"")' % (fila, fila, fila, fila))
        else:
            motor.val(ws, 'G%d' % fila, 'No aplica')
            motor.val(ws, 'H%d' % fila, '')
        motor.f(ws, 'I%d' % fila,
                '=IF(H%d="No",IF(C%d="Sí","DESCARTAR","Punto débil"),'
                'IF(E%d="","",IF(E%d="No",IF(C%d="Sí","DESCARTAR",'
                '"Punto débil"),IF(E%d="Por comprobar","Pendiente","OK"))))'
                % (fila, fila, fila, fila, fila, fila))
        motor.val(ws, 'J%d' % fila, notatxt, wrap=True)
        if idlegal:
            nota_legal_celda(ws, 'J%d' % fila, idlegal)
        ws.row_dimensions[fila].height = 58
        fila += 1
    l_fin = fila - 1
    # Hallazgo B2 (2026-09-10): esta celda llevaba la nota de PA-07, y el
    # `dato` de PA-07 dice justo lo CONTRARIO de lo que sugiere citarlo aquí
    # («el DB-HS 3 NO se aplica a un obrador: remite al RITE»). `nota_legal()`
    # sólo compone «Verificado · norma · URL», nunca el `dato`, así que un
    # hallazgo NEGATIVO se leía como una verificación positiva del CTE. La
    # base correcta para esta casilla es PA-07b (RITE), la misma que ya cita
    # el ítem siguiente.
    nota_legal_celda(ws, 'D%d' % L_INI, 'PA-07b')
    motor.semaforo_texto(ws, 'I%d:I%d' % (L_INI, l_fin),
                         (('DESCARTAR', motor.CF_ROJO_BG, motor.CF_ROJO_FG),
                          ('Punto débil', motor.CF_AMBAR_BG,
                           motor.CF_AMBAR_FG),
                          ('Pendiente', motor.CF_GRIS_BG, motor.CF_GRIS_FG),
                          ('OK', motor.CF_VERDE_BG, motor.CF_VERDE_FG)))
    motor.semaforo_texto(ws, 'H%d:H%d' % (L_INI, l_fin),
                         (('No', motor.CF_ROJO_BG, motor.CF_ROJO_FG),
                          ('Sí', motor.CF_VERDE_BG, motor.CF_VERDE_FG)))

    f = l_fin + 2
    filas_res = [
        ('Ítems de la ficha', '=COUNTIF(C%d:C%d,"Sí")+COUNTIF(C%d:C%d,"No")'
         % (L_INI, l_fin, L_INI, l_fin), C.FMT_ENT),
        ('De los cuales, eliminatorios', '=COUNTIF(C%d:C%d,"Sí")'
         % (L_INI, l_fin), C.FMT_ENT),
        ('Motivos de descarte', '=COUNTIF(I%d:I%d,"DESCARTAR")'
         % (L_INI, l_fin), C.FMT_ENT),
        ('Puntos débiles', '=COUNTIF(I%d:I%d,"Punto débil")' % (L_INI, l_fin),
         C.FMT_ENT),
        ('Pendientes de comprobar', '=COUNTIF(I%d:I%d,"Pendiente")'
         % (L_INI, l_fin), C.FMT_ENT),
        ('Resueltos sin problema', '=COUNTIF(I%d:I%d,"OK")' % (L_INI, l_fin),
         C.FMT_ENT),
    ]
    for k, (etq, formula, fmt) in enumerate(filas_res):
        motor.val(ws, 'B%d' % (f + k), etq)
        motor.f(ws, 'E%d' % (f + k), formula, fmt=fmt)
    f_tot, f_desc_n, f_pend = f, f + 2, f + 4
    f_avance = f + len(filas_res)
    motor.val(ws, 'B%d' % f_avance, 'Ficha comprobada')
    motor.f(ws, 'E%d' % f_avance,
            motor.iferror('(E%d-E%d)/E%d' % (f_tot, f_pend, f_tot)),
            fmt=C.FMT_PCT)
    f_ver = f_avance + 1
    motor.val(ws, 'B%d' % f_ver, 'VEREDICTO DEL LOCAL', bold=True)
    motor.f(ws, 'E%d' % f_ver,
            '=IF(E%d>0,"DESCARTAR ESTE LOCAL",IF(E%d>0,'
            '"NO DECIDAS TODAVÍA",IF(E%d>0,"SIGUE ADELANTE CON RESERVAS",'
            '"SIGUE ADELANTE")))' % (f_desc_n, f_pend, f + 3), bold=True)
    C.destacado(ws, 'E%d' % f_ver)
    motor.semaforo_texto(ws, 'E%d' % f_ver,
                         (('DESCARTAR ESTE LOCAL', motor.CF_ROJO_BG,
                           motor.CF_ROJO_FG),
                          ('NO DECIDAS TODAVÍA', motor.CF_GRIS_BG,
                           motor.CF_GRIS_FG),
                          ('SIGUE ADELANTE CON RESERVAS', motor.CF_AMBAR_BG,
                           motor.CF_AMBAR_FG),
                          ('SIGUE ADELANTE', motor.CF_VERDE_BG,
                           motor.CF_VERDE_FG)))

    fila = f_ver + 2
    refs, fila = C.bloque_listas(ws, fila, [('Respuesta', RESPUESTAS)],
                                 col='B')
    C.dv_rango(ws, ['E%d' % r for r in range(L_INI, l_fin + 1)],
               refs['Respuesta'], 'Respuesta',
               'Elige «Sí», «No» o «Por comprobar» de la lista.')
    C.parrafo(ws, fila,
              'Los ocho ítems eliminatorios no lo son por gusto: son los '
              'defectos que no se arreglan con dinero razonable, o que no '
              'dependen de ti (la comunidad, el planeamiento, la acometida, la '
              'cubierta). El resto encarecen la obra y hay que llevarlos al '
              'libro 2 como partida, no descartar el local por ellos. Y guarda '
              'la ficha de los locales que descartes: comparar tres fichas es '
              'lo que enseña qué estás aceptando sin darte cuenta.',
              col_ini='B', col_fin='J', alto=56)
    C.pagina(ws, apaisado=True, titulos='5:5')
    return ws


# --------------------------------------------------------------------------
# Mapa de celdas citables
# --------------------------------------------------------------------------
def mapa_celdas():
    """`{etiqueta: {ref, tipo}}` — el `valor` lo rellena la verificación
    `data_only` después de `inject_cache`."""
    n_fin = C_INI + len(EQUIPOS_CAPACIDAD) - 1
    p_fin = K_PICO_INI + len(D.PICOS) - 1
    l_fin = L_INI + len(VISITA) - 1
    f_res = l_fin + 2
    m = {}

    def add(etq, hoja, celda, tipo):
        m[etq] = {'ref': '%s.xlsx!%s!%s' % (NOMBRE, hoja, celda),
                  'tipo': tipo}

    # --- Parámetros ------------------------------------------------------
    add('Horas de turno del obrador', H_PAR, 'B%d' % fila_par(0), 'entrada')
    add('Horas productivas sobre la jornada', H_PAR, 'B%d' % fila_par(2),
        'entrada')
    add('Minutos disponibles al día', H_PAR, 'B%d' % fila_par(P_MIN_DIA),
        'salida')
    add('Días de apertura al año', H_PAR, 'B%d' % fila_par(5), 'entrada')
    add('Piezas al día en velocidad de crucero', H_PAR, 'B%d' % fila_par(6),
        'entrada')
    add('Margen de seguridad exigido', H_PAR, 'B%d' % fila_par(7), 'parametro')
    add('Superficie total del local', H_PAR, 'B%d' % fila_par(8), 'entrada')
    add('Parte mínima del local dedicada a obrador', H_PAR,
        'B%d' % fila_par(9), 'parametro')
    add('Potencia de los hornos', H_PAR, 'B%d' % fila_par(10), 'entrada')
    add('Potencia instalada total', H_PAR, 'B%d' % fila_par(11), 'entrada')
    add('Caudal de campana necesario', H_PAR, 'B%d' % fila_par(P_CAUDAL),
        'salida')
    add('Altura libre mínima exigida', H_PAR, 'B%d' % fila_par(16), 'entrada')
    add('Número de zonas del local', H_PAR, 'B%d' % fila_par(17), 'parametro')

    # --- Zonas -----------------------------------------------------------
    add('Metros cuadrados repartidos entre las zonas', H_ZON, 'D14', 'salida')
    add('Descuadre de metros contra la superficie declarada', H_ZON, 'D16',
        'salida')
    add('Metros de obrador', H_ZON, 'D18', 'salida')
    add('Parte del local dedicada a obrador', H_ZON, 'D19', 'salida')
    add('Metros de zona de venta', H_ZON, 'D20', 'salida')
    add('Parte del local dedicada a venta', H_ZON, 'D21', 'salida')
    add('Veredicto del reparto de metros', H_ZON, 'D23', 'salida')
    add('Numeración del recorrido correcta', H_ZON, 'D28', 'salida')
    add('Veredicto de marcha adelante', H_ZON, 'D35', 'salida')
    add('Metros cuadrados del obrador (celda de entrada)', H_ZON,
        'D%d' % Z_FILA['Obrador'], 'entrada')

    # --- Capacidad -------------------------------------------------------
    for i, eq in enumerate(EQUIPOS_CAPACIDAD):
        add('Piezas/día de obrador que permite: ' + eq[0], H_CAP,
            'M%d' % (C_INI + i), 'salida')
    add('Equipos que dices tener', H_CAP, 'E%d' % (n_fin + 3), 'salida')
    add('Minutos disponibles al día (capacidad)', H_CAP, 'E%d' % (n_fin + 2),
        'salida')
    add('Piezas por ciclo del abatidor', H_CAP, 'H%d' % (C_INI + 6), 'salida')
    add('Ciclos al día del abatidor', H_CAP, 'J%d' % (C_INI + 6), 'salida')
    add('Parte del surtido que pasa por el abatidor', H_CAP,
        'L%d' % (C_INI + 6), 'entrada')

    # --- Cuello de botella ------------------------------------------------
    add('Piezas al día objetivo', H_CUE, 'B%d' % K_OBJ, 'salida')
    add('Piezas al día que permite el conjunto', H_CUE, 'B%d' % K_CAP,
        'salida')
    add('El equipo que limita', H_CUE, 'B%d' % K_EQU, 'salida')
    add('Holgura sobre el día normal (piezas)', H_CUE, 'B%d' % K_HOL, 'salida')
    add('Holgura sobre el día normal (%)', H_CUE, 'B%d' % K_HOLP, 'salida')
    add('Veredicto del día normal', H_CUE, 'B%d' % K_VER, 'salida')
    add('El segundo equipo más justo permite', H_CUE, 'B%d' % K_SEG, 'salida')
    add('El segundo equipo más justo es', H_CUE, 'B%d' % K_SEGE, 'salida')
    for i, p in enumerate(D.PICOS):
        add('Piezas/día que hacen falta en ' + p['nombre'], H_CUE,
            'D%d' % (K_PICO_INI + i), 'salida')
        add('Déficit de capacidad en ' + p['nombre'], H_CUE,
            'F%d' % (K_PICO_INI + i), 'salida')
    add('Picos que no aguantas', H_CUE, 'B%d' % (p_fin + 2), 'salida')
    add('El peor déficit de campaña', H_CUE, 'B%d' % (p_fin + 3), 'salida')
    add('Pico del peor déficit', H_CUE, 'B%d' % (p_fin + 4), 'salida')
    add('Veredicto de campaña', H_CUE, 'B%d' % (p_fin + 5), 'salida')

    # --- Ficha de visita --------------------------------------------------
    add('Ítems de la ficha de visita', H_LOC, 'E%d' % f_res, 'salida')
    add('Ítems eliminatorios de la ficha', H_LOC, 'E%d' % (f_res + 1),
        'salida')
    add('Motivos de descarte del local', H_LOC, 'E%d' % (f_res + 2), 'salida')
    add('Puntos débiles del local', H_LOC, 'E%d' % (f_res + 3), 'salida')
    add('Pendientes de comprobar del local', H_LOC, 'E%d' % (f_res + 4),
        'salida')
    add('Ficha comprobada (%)', H_LOC, 'E%d' % (f_res + 6), 'salida')
    add('VEREDICTO DEL LOCAL', H_LOC, 'E%d' % (f_res + 7), 'salida')
    add('Caudal de campana ofrecido por el local', H_LOC, 'F%d' % (L_INI + 2),
        'entrada')
    add('Potencia contratable del local', H_LOC, 'F%d' % (L_INI + 3),
        'entrada')
    add('Altura libre medida en el local', H_LOC, 'F%d' % (L_INI + 6),
        'entrada')
    add('Superficie útil medida en el local', H_LOC, 'F%d' % (L_INI + 7),
        'entrada')
    return m


# --------------------------------------------------------------------------
def construir():
    D.gate_legal(IDS_LEGALES)
    wb = Workbook()
    wb.remove(wb.active)
    hoja_instrucciones(wb)
    hoja_parametros(wb)
    hoja_zonas(wb)
    hoja_capacidad(wb)
    hoja_cuello(wb)
    hoja_local(wb)

    wb.properties.creator = 'AI Chef Pro'
    wb.properties.lastModifiedBy = 'AI Chef Pro'
    wb.properties.title = TITULO
    wb.properties.subject = C.PRODUCTO + ' · Versión 1.0 · septiembre 2026'

    verdes = {}
    for ws in wb.worksheets:
        motor.retirar_verde_de_calculadas(ws)
        verdes[ws.title] = motor.proteger(ws)
    wb.calculation.fullCalcOnLoad = True

    destino = os.path.join(AQUI, 'build')
    if not os.path.isdir(destino):
        os.makedirs(destino)
    ruta = os.path.join(destino, NOMBRE + '.xlsx')
    wb.save(ruta)
    C.barrer_cp1252(wb)
    return ruta, verdes


def gate_totales(ruta):
    """`=SUM(rango)` recalculado a mano contra el caché (hallazgos A1/A3,
    2026-09-10): `pycel` cachea como 0 un `SUM` sobre una columna de
    `SUMPRODUCT`. Se llama DESPUÉS de `inject_cache.py`."""
    wbv = openpyxl.load_workbook(ruta, data_only=True)
    incoherentes = C.gate_sum_rango(wbv)
    wbv.close()
    if incoherentes:
        raise SystemExit('TOTALES con caché incoherente:\n  '
                         + '\n  '.join(incoherentes))


def verificar_y_mapear(ruta):
    """`inject_cache` + comprobación `data_only` de CADA fórmula registrada +
    `mapa-<libro>.json` con el valor cacheado.

    Ojo con las fórmulas que devuelven `""`: openpyxl con `data_only=True`
    las lee como `None`, exactamente igual que una fórmula que se quedó SIN
    cachear. Distinguirlas a ojo es imposible, así que la comprobación pasa por
    pycel: sólo se acepta el `None` cuando pycel dice que esa fórmula vale
    cadena vacía -que es el «sin dato» de la familia-.
    """
    import logging
    logging.disable(logging.CRITICAL)
    from pycel import ExcelCompiler
    subprocess.check_call(['/usr/local/bin/python3',
                           os.path.join(RAIZ, 'inject_cache.py'), ruta])
    wb = openpyxl.load_workbook(ruta, data_only=True)
    calc = ExcelCompiler(ruta)
    fallos, vacias = [], 0
    for hoja, coord, formula in motor.REGISTRO:
        v = wb[hoja][coord].value
        if isinstance(v, str) and v.startswith('#'):
            fallos.append('%s!%s = %s (%s)' % (hoja, coord, v, formula[:60]))
            continue
        if v is not None:
            continue
        esperado = calc.evaluate("'%s'!%s" % (hoja, coord))
        if esperado == '':
            vacias += 1                     # «sin dato» legítimo, no un hueco
        else:
            fallos.append('%s!%s SIN VALOR, pycel dice %r (%s)'
                          % (hoja, coord, esperado, formula[:60]))
    m = mapa_celdas()
    for etq, d in m.items():
        _f, hoja, celda = d['ref'].split('!')
        v = wb[hoja][celda].value
        if v is None and wb[hoja][celda].data_type != 'n':
            v = calc.evaluate("'%s'!%s" % (hoja, celda))
        d['valor'] = v
    destino = os.path.dirname(ruta)
    with open(os.path.join(destino, 'mapa-' + NOMBRE + '.json'), 'w',
              encoding='utf-8') as fh:
        json.dump(m, fh, ensure_ascii=False, indent=1)
    return fallos, m, vacias


# --------------------------------------------------------------------------
def demo(ruta):
    """Seis comportamientos que el libro tiene que tener, probados con pycel
    sobre el fichero que se acaba de escribir.

    Regla de pycel aprendida aquí: hay que **evaluar la salida ANTES de tocar
    la entrada**. `set_value()` invalida los nodos que ya están en el grafo, y
    los que no están todavía se construyen leyendo el valor CACHEADO del
    fichero: si se toca la entrada primero, la salida sale con el número viejo
    y la prueba pasa por casualidad.
    """
    import logging
    logging.disable(logging.CRITICAL)
    from pycel import ExcelCompiler
    resultados = []

    def prueba(nombre, ok, detalle=''):
        resultados.append((nombre, bool(ok), detalle))

    def compilador(salidas):
        c = ExcelCompiler(ruta)
        for a in salidas:
            c.evaluate(a)
        return c

    A_CAP = "'%s'!B%d" % (H_CUE, K_CAP)
    A_EQU = "'%s'!B%d" % (H_CUE, K_EQU)
    A_VER = "'%s'!B%d" % (H_CUE, K_VER)
    l_fin = L_INI + len(VISITA) - 1
    A_LOC = "'%s'!E%d" % (H_LOC, l_fin + 2 + 7)

    # 1. El cuello de botella se identifica y el veredicto es coherente.
    c = compilador([A_CAP, A_EQU, A_VER])
    cap, equ, ver = c.evaluate(A_CAP), c.evaluate(A_EQU), c.evaluate(A_VER)
    prueba('El cuello de botella se identifica y trae nombre de equipo',
           isinstance(cap, (int, float)) and isinstance(equ, str) and equ
           and ver in ('NO LLEGAS', 'AJUSTADO', 'SUFICIENTE'),
           'capacidad=%s · equipo=%s · veredicto=%s' % (cap, equ, ver))

    # 2. Quitar el equipo que limita mueve el cuello a otro equipo.
    c = compilador([A_CAP, A_EQU])
    c.set_value("'%s'!D%d" % (H_CAP, C_INI + 6), 'No')
    cap2, equ2 = c.evaluate(A_CAP), c.evaluate(A_EQU)
    prueba('Quitar el equipo que limita mueve el cuello a otro equipo',
           equ2 != equ and cap2 > cap,
           'ahora limita %s con %s piezas/día' % (equ2, round(cap2)))

    # 3. Doblar el turno cambia el veredicto del día normal.
    c = compilador([A_CAP, A_VER])
    c.set_value("'%s'!B%d" % (H_PAR, fila_par(1)), 2)
    cap3, ver3 = c.evaluate(A_CAP), c.evaluate(A_VER)
    prueba('Con dos turnos el veredicto del día normal pasa a SUFICIENTE',
           ver3 == 'SUFICIENTE' and cap3 > cap,
           'veredicto=%s · capacidad=%s' % (ver3, round(cap3)))

    # 4. Un «No» en un ítem eliminatorio manda DESCARTAR.
    c = compilador([A_LOC])
    base = c.evaluate(A_LOC)
    c.set_value("'%s'!E%d" % (H_LOC, L_INI), 'No')       # ítem 1, eliminatorio
    tras = c.evaluate(A_LOC)
    prueba('Un «No» en un ítem eliminatorio manda DESCARTAR',
           tras == 'DESCARTAR ESTE LOCAL' and base != tras,
           'antes=%s · después=%s' % (base, tras))

    # 5. Un «No» en un ítem NO eliminatorio no descarta el local.
    c = compilador([A_LOC])
    c.set_value("'%s'!E%d" % (H_LOC, L_INI + 6), 'No')   # altura libre
    tras5 = c.evaluate(A_LOC)
    prueba('Un «No» en un ítem no eliminatorio NO descarta el local',
           tras5 != 'DESCARTAR ESTE LOCAL', 'veredicto=%s' % tras5)

    # 6. El semáforo del dato medido no se enciende con texto.
    A_H = "'%s'!H%d" % (H_LOC, L_INI + 2)
    A_I = "'%s'!I%d" % (H_LOC, L_INI + 2)
    c = compilador([A_H, A_I])
    c.set_value("'%s'!F%d" % (H_LOC, L_INI + 2), 'por pedir')
    cumple, vered = c.evaluate(A_H), c.evaluate(A_I)
    prueba('Un texto en el dato medido no enciende el semáforo: deja ""',
           cumple == '' and vered == 'OK',
           '¿cumple?=%r · veredicto=%r' % (cumple, vered))

    # 7. Los metros de las zonas cuadran con la superficie declarada.
    c = ExcelCompiler(ruta)
    desc = c.evaluate("'%s'!D16" % H_ZON)
    tot = c.evaluate("'%s'!D14" % H_ZON)
    prueba('El reparto de metros suma la superficie declarada',
           abs(desc) < 0.001 and abs(tot - D.NEGOCIO['m2_total']) < 0.001,
           'total=%s · descuadre=%s' % (tot, desc))
    return resultados


# --------------------------------------------------------------------------
def main():
    ruta, verdes = construir()
    print('escrito:', ruta)
    fallos, m, vacias = verificar_y_mapear(ruta)
    gate_totales(ruta)
    res = demo(ruta)
    n_verdes = sum(verdes.values())
    print('-' * 66)
    print('fórmulas registradas : %d' % len(motor.REGISTRO))
    print('fórmulas sin valor   : %d' % len(fallos))
    print('fórmulas que dan ""  : %d  (sin dato legítimo)' % vacias)
    for x in fallos[:20]:
        print('   ', x)
    print('celdas verdes        : %d  %s' % (n_verdes, verdes))
    print('notas legales        : %d' % NOTAS_LEGALES['n'])
    print('etiquetas del mapa   : %d' % len(m))
    print('demos:')
    for nombre, ok, detalle in res:
        print('   [%s] %s — %s' % ('OK' if ok else 'FALLA', nombre, detalle))
    todo_ok = not fallos and all(ok for _n, ok, _d in res)
    print('-' * 66)
    print('RESULTADO:', 'VERDE' if todo_ok else 'ROJO')
    return 0 if todo_ok else 1


if __name__ == '__main__':
    sys.exit(main())
