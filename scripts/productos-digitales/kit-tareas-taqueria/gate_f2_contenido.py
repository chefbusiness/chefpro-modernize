#!/usr/bin/env python3
"""gate_f2_contenido.py — Ronda 2 de la F2 de la Taquería = GATE DE SCRIPT (SPEC §5). Mide los xlsx generados contra la
SPEC: rangos de tareas por hoja (§3.1), molde (fila 4, DV, 5 filas libres, contador, pie, protección), umbrales (63/74 °C,
nunca 65/75), registro horario, Cinco de Mayo, tildes, glosario, prohibiciones, no latinos, U+202F, propiedades del libro.
Uso: /usr/local/bin/python3 gate_f2_contenido.py [dir_kit] [--spec ruta]. Exit 0 = PASS."""
import sys, re, unicodedata, pathlib, glob
from openpyxl import load_workbook
args = [a for a in sys.argv[1:] if not a.startswith('--')]
KIT = pathlib.Path(args[0] if args else 'astro-site/public/dl/kit-tareas-taqueria')
SPEC = pathlib.Path(sys.argv[sys.argv.index('--spec') + 1] if '--spec' in sys.argv else 'scripts/productos-digitales/kit-tareas-taqueria/02-SPEC-kit-tareas-taqueria.md')
PIE = '— Kit de Tareas Recurrentes · Taquería Mexicana · AI Chef Pro · aichef.pro'
DV_OK = '"✓,—,N/A"'
VERDE = '00E8F5E9'
fallos, avisos = [], []
F = lambda m: fallos.append(m)
# --- rangos esperados por fichero (SPEC §3.1) ---
spec = SPEC.read_text(encoding='utf-8')
m = re.search(r'#+\s*§?\s*3\.1[^\n]*\n(.*?)(?=\n#{2,3} )', spec, re.S)
esperado = {}   # prefijo de fichero -> lista de (hoja o None, min, max)
fich = None
for fila in m.group(1).split('\n'):
    if not re.match(r'\|\s*0\d\s*\|', fila): continue
    celdas = [c.strip() for c in fila.strip().strip('|').split('|')]
    num, nombre, hoja, obj = celdas[0], celdas[1].strip('`'), celdas[2], celdas[3]
    if nombre: fich = nombre
    r = re.findall(r'(\d{1,3})\s*[-–]\s*(\d{1,3})', obj)
    if not r: continue
    if 'cada una' in obj:            # 06: 6 hojas × 7-9
        a, b = map(int, r[0]); esperado.setdefault(fich, []).extend([(None, a, b)] * 6)
    else:
        a, b = map(int, r[-1]); esperado.setdefault(fich, []).append((hoja, a, b))
# --- recorrido de los xlsx ---
ficheros = sorted(glob.glob(str(KIT / '*.xlsx')))
if len(ficheros) != 11: F(f'{len(ficheros)} xlsx en {KIT}, esperados 11')
total = 0
textos_por_fichero = {}
PROHIBIDO = [(r'\b65\s*°\s*C', '65 °C derogado'), (r'\b75\s*°\s*C', '75 °C derogado'), (r'registro horario digital|control horario digital|fichaje digital|horario digital|obligatorio en 2026|digital obligatori', 'registro horario digital'),
             (r'tex-?mex', 'tex-mex'), (r'\bgyro', 'gyro'), (r'\bburrito', 'burrito'), (r'\bnachos', 'nachos'), (r'\bpesos\b|\bMXN\b|\$', 'importes en pesos'), (r' |‑', 'U+202F/U+2011')]
SIN_TILDE = re.compile(r'\b(camara|limite|preparacion|maria|numero|sesamo|rotacion|alergeno|alergenos|dia|dias|aqui|segun|tambien|hora limite|limpieza basica|trazabilidad y limite)\b', re.I)
GLOSAS = [('trompo', r'trompo \(asador vertical\)'), ('tortillería', r'tortillería \(obrador'), ('guisado', r'guisados? \(cazuelas?'), ('APPCC', r'APPCC \(HACCP\)'), ('escandallo', r'escandallo \(costeo\)')]
for path in ficheros:
    nombre = pathlib.Path(path).name
    wb = load_workbook(path)
    pr = wb.properties
    if not (pr.subject and 'Taquería Mexicana' in pr.subject and 'v2.0' in pr.subject): F(f'{nombre}: subject «{pr.subject}»')
    if pr.creator != 'AI Chef Pro': F(f'{nombre}: creator «{pr.creator}»')
    if wb.worksheets[0].title != 'Instrucciones': F(f'{nombre}: la primera hoja no es Instrucciones')
    textos = []
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str): textos.append((ws.title, c.coordinate, c.value))
    textos_por_fichero[nombre] = textos
    ins = '\n'.join(v for t, _, v in textos if t == 'Instrucciones')
    es_bonus = nombre.startswith('BONUS')
    for req in ('Cómo usar', 'Diseñado por John Guerrero', 'Versión 2.0', PIE) + (() if es_bonus else ('N/A',)):
        if req not in ins: F(f'{nombre}/Instrucciones: falta «{req[:40]}»')
    if 'Se conecta con' not in ins and 'Dónde encaja' not in ins: F(f'{nombre}/Instrucciones: sin bloque de conexión')
    # hojas de checklist
    prefijo = next((k for k in esperado if nombre.startswith(k.split('.')[0][:2])), None)
    rangos = list(esperado.get(prefijo, [])) if prefijo else []
    hojas_check = []
    for ws in wb.worksheets[1:]:
        a4, b4 = ws['A4'].value, ws['B4'].value
        es_check = (a4 in ('Nº', '#')) and b4 == 'Tarea'
        ultimo = max((r for r in range(1, ws.max_row + 1) if any(ws.cell(r, c).value is not None for c in range(1, 8))), default=0)
        pie = ' '.join(str(ws.cell(ultimo, c).value) for c in range(1, 8) if ws.cell(ultimo, c).value)
        if not pie.startswith(PIE): F(f'{nombre}/{ws.title}: pie «{pie[:50]}»')
        if not ws.protection.sheet: F(f'{nombre}/{ws.title}: sin protección')
        if not ws.print_area: F(f'{nombre}/{ws.title}: sin print_area')
        if not es_check: continue
        if a4 != 'Nº' or 'º' not in a4: F(f'{nombre}/{ws.title}: A4 debe ser N + U+00BA')
        if ws.freeze_panes != 'A5': F(f'{nombre}/{ws.title}: freeze {ws.freeze_panes}')
        if [ws.cell(4, c).value for c in range(2, 8)] != ['Tarea', 'Zona', 'Responsable', ws['E4'].value, '✓ Completada', 'Firma']: F(f'{nombre}/{ws.title}: cabecera {[ws.cell(4, c).value for c in range(1, 8)]}')
        dv_celdas = set()
        for dv in ws.data_validations.dataValidation:
            if dv.formula1 != DV_OK: F(f'{nombre}/{ws.title}: DV {dv.formula1}')
            for rng in str(dv.sqref).split():
                if ':' in rng:
                    c0, c1 = rng.split(':'); r0, r1 = int(re.sub(r'\D', '', c0)), int(re.sub(r'\D', '', c1))
                    dv_celdas.update(f'F{r}' for r in range(r0, r1 + 1))
                else: dv_celdas.add(rng)
        tareas, vacias_num, libres, contador = 0, 0, 0, None
        for r in range(5, ws.max_row + 1):
            a, b = ws.cell(r, 1).value, ws.cell(r, 2).value
            if isinstance(a, str) and a.startswith('Tareas completadas'): contador = r; break
            if isinstance(a, (int, float)):
                if b: 
                    tareas += 1
                    if not ws.cell(r, 4).value or not ws.cell(r, 5).value: F(f'{nombre}/{ws.title}:{r} Responsable/Hora vacíos (B6)')
                    if f'F{r}' not in dv_celdas: F(f'{nombre}/{ws.title}:F{r} sin desplegable')
                    base_txt = re.sub(r' \(si tienes el Pack APPCC[^)]*\)| — anota la lectura: ____ °C| \(refrigeración 0-4 °C\)', '', str(b))
                    if len(base_txt) > 110: F(f'{nombre}/{ws.title}:B{r} > 110 caracteres ({len(base_txt)})')
                else: vacias_num += 1
            elif a is None and b is None and ws.cell(r, 1).fill.fill_type == 'solid' and ws.cell(r, 1).fill.fgColor.rgb == VERDE:
                libres += 1
                if f'F{r}' not in dv_celdas: F(f'{nombre}/{ws.title}:F{r} fila libre sin desplegable')
        if contador is None: F(f'{nombre}/{ws.title}: sin fila «Tareas completadas»')
        else:
            fd, ff = str(ws.cell(contador, 4).value), str(ws.cell(contador, 6).value)
            if not (fd.startswith('=COUNTIFS(B5:B') and '"?*"' in fd and '"✓"' in fd): F(f'{nombre}/{ws.title}: contador D «{fd}»')
            if not (ff.startswith('=COUNTIF(B5:B') and '"N/A"' in ff): F(f'{nombre}/{ws.title}: contador F «{ff}»')
        es_plantilla = nombre.startswith('09')
        if vacias_num and not es_plantilla: F(f'{nombre}/{ws.title}: {vacias_num} filas numeradas sin tarea')
        if es_plantilla and vacias_num != 15: F(f'{nombre}/{ws.title}: {vacias_num} filas numeradas en blanco (el molde base lleva 15)')
        if not es_plantilla and libres != 5: F(f'{nombre}/{ws.title}: {libres} filas libres verdes (esperadas 5)')
        if not es_plantilla: hojas_check.append((ws.title, tareas)); total += tareas
    # rangos por hoja
    if rangos:
        if len(hojas_check) != len(rangos): F(f'{nombre}: {len(hojas_check)} hojas de checklist, esperadas {len(rangos)}')
        for (titulo, n), (hoja, a, b) in zip(hojas_check, rangos):
            if hoja and hoja != titulo: avisos.append(f'{nombre}: hoja «{titulo}» ≠ SPEC «{hoja}»')
            if not (a <= n <= b): F(f'{nombre}/{titulo}: {n} tareas, rango SPEC {a}-{b}')
    # texto: prohibiciones, tildes, no latinos, glosario, Cinco de Mayo
    for t, coord, v in textos:
        for rx, nom in PROHIBIDO:
            if re.search(rx, v): F(f'{nombre}/{t}!{coord}: {nom}: «{v[:60]}»')
        if any(unicodedata.category(ch).startswith('L') and any(k in unicodedata.name(ch, '') for k in ('CJK', 'CYRILLIC', 'HANGUL', 'ARABIC', 'HEBREW', 'THAI', 'KATAKANA', 'HIRAGANA')) for ch in v): F(f'{nombre}/{t}!{coord}: no latino')
        mt = SIN_TILDE.search(re.sub(r'\S+\.xlsx', '', v))
        if mt and not re.search(r'johnguerrero|aichef|@', v): F(f'{nombre}/{t}!{coord}: sin tilde «{mt.group(0)}» en «{v[:50]}»')
        if 'Cinco de Mayo' in v and not re.search(r'EE\. ?UU\.|estadounidense|Puebla', v): F(f'{nombre}/{t}!{coord}: Cinco de Mayo sin matiz')
    todo = '\n'.join(v for t, coord, v in textos if t != 'Instrucciones' and (nombre.startswith('BONUS') or coord.startswith('B')))
    for palabra, glosa in GLOSAS:
        if re.search(r'\b' + palabra + r'\b', todo) and not re.search(glosa, todo): F(f'{nombre}: usa «{palabra}» sin la glosa «{glosa}»')   # minúsculas = prosa; los rótulos en MAYÚSCULAS no cuentan
print(f'ficheros {len(ficheros)} · tareas numeradas escritas {total} (SPEC 270-330) · avisos {len(avisos)}')
for a in avisos: print('  aviso:', a)
if not (270 <= total <= 330) and ficheros: F(f'total {total} fuera de 270-330')
if fallos:
    print(f'GATE F2 CONTENIDO: FAIL ({len(fallos)})'); [print(' -', f) for f in fallos[:80]]; sys.exit(1)
print('GATE F2 CONTENIDO: PASS')
