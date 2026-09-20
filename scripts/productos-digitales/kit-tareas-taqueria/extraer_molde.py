#!/usr/bin/env python3
"""extraer_molde.py — Volcado ESTRUCTURAL (no de contenido) de las hojas de referencia de la familia Kit de Tareas,
para que el contrato de helpers de F2 se escriba contra hechos medidos y para comparar después los xlsx generados.
Uso: /usr/local/bin/python3 extraer_molde.py [xlsx ...]  → JSON por stdout. Ligero: abre pocos ficheros, uno a uno."""
import sys, json
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
def hoja(ws):
    d = {'title': ws.title, 'max_row': ws.max_row, 'max_col': ws.max_column,
         'freeze': str(ws.freeze_panes), 'print_area': ws.print_area, 'protected': ws.protection.sheet,
         'widths': {k: round(v.width, 2) for k, v in ws.column_dimensions.items() if v.width},
         'merged': [str(r) for r in ws.merged_cells.ranges][:12],
         'dv': [{'sqref': str(dv.sqref), 'type': dv.type, 'formula1': dv.formula1, 'allow_blank': dv.allow_blank} for dv in ws.data_validations.dataValidation],
         'cf': [str(r.sqref) for r in ws.conditional_formatting][:8],
         'rows': []}
    for r in range(1, min(ws.max_row, 9) + 1):
        fila = []
        for c in range(1, min(ws.max_column, 8) + 1):
            cell = ws.cell(r, c)
            v = cell.value
            if v is None and not cell.fill.fgColor.rgb not in (None, '00000000'): continue
            fila.append({'c': get_column_letter(c), 'v': (str(v)[:60] if v is not None else None),
                         'bold': bool(cell.font.bold), 'sz': cell.font.sz, 'color': cell.font.color.rgb if cell.font.color and cell.font.color.type == 'rgb' else None,
                         'fill': cell.fill.fgColor.rgb if cell.fill and cell.fill.fill_type == 'solid' else None,
                         'align': cell.alignment.horizontal, 'wrap': bool(cell.alignment.wrap_text), 'locked': cell.protection.locked})
        d['rows'].append({'r': r, 'h': ws.row_dimensions[r].height, 'cells': [x for x in fila if x['v'] is not None or x['fill']]})
    # últimas 8 filas (contador, filas libres, pie)
    d['tail'] = []
    for r in range(max(1, ws.max_row - 7), ws.max_row + 1):
        vals = [(get_column_letter(c), str(ws.cell(r, c).value)[:70], ws.cell(r, c).fill.fgColor.rgb if ws.cell(r, c).fill.fill_type == 'solid' else None)
                for c in range(1, min(ws.max_column, 8) + 1) if ws.cell(r, c).value is not None or ws.cell(r, c).fill.fill_type == 'solid']
        d['tail'].append({'r': r, 'cells': vals})
    return d
out = {}
for path in sys.argv[1:]:
    wb = load_workbook(path)
    out[path.split('/')[-2] + '/' + path.split('/')[-1]] = {'sheets': [ws.title for ws in wb.worksheets], 'detalle': [hoja(ws) for ws in wb.worksheets[:3]],
                                                             'props': {'title': wb.properties.title, 'subject': wb.properties.subject, 'creator': wb.properties.creator}}
print(json.dumps(out, ensure_ascii=False, indent=1))
