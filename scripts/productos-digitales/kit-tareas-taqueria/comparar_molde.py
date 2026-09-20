#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""comparar_molde.py — ¿la hoja de checklist que emite el generador de la taquería es el MISMO
molde estructural que la del kit base v2.0?

Compara la hoja 2 de un xlsx generado contra `astro-site/public/dl/kit-tareas/01-apertura-cierre.xlsx`
hoja «Apertura Cocina» (la referencia canónica del contrato `03-contrato-molde-v2.md`):

  · `freeze_panes` y anchos de columna
  · textos + fills + bold de las filas 1-4 (el texto propio del título y de la fila 2 se compara
    por FORMA, no por literal: son del kit, no del molde)
  · `formula1` / `allow_blank` de la validación de datos
  · patrón de las fórmulas del contador (numerador COUNTIFS y denominador honesto)
  · pie, `protected`, `print_area` no vacío
  · 5 filas libres verdes SIN número justo antes de la fila del contador

Uso:
    /usr/local/bin/python3 comparar_molde.py <xlsx generado> [hoja]
"""
import os
import re
import sys

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

AQUI = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, AQUI)
# `extraer_molde.py` es un script sin guarda `__main__`: al importarlo lee `sys.argv` y escupe su
# JSON por stdout. Se importa con los argumentos vaciados y la salida silenciada; lo que se
# reutiliza es su función `hoja()`, que es la que define QUÉ se mide.
_argv, _stdout = sys.argv, sys.stdout
try:
    sys.argv = [_argv[0]]
    sys.stdout = open(os.devnull, 'w')
    from extraer_molde import hoja                               # noqa: E402
finally:
    sys.stdout.close()
    sys.argv, sys.stdout = _argv, _stdout

RAIZ = os.path.dirname(os.path.dirname(os.path.dirname(AQUI)))
REF = os.path.join(RAIZ, 'astro-site', 'public', 'dl', 'kit-tareas',
                   '01-apertura-cierre.xlsx')
REF_HOJA = 'Apertura Cocina'

VERDE = 'E8F5E9'
HOLGURA = 5
#: Los rótulos de la fila 4 que SON del molde. «Hora Límite» admite las variantes que el motor
#: impone por contenido (`motor.cadencia`): lo que no puede cambiar es el resto.
CAB_FIJA = {1: 'Nº', 2: 'Tarea', 3: 'Zona', 4: 'Responsable',
            6: '✓ Completada', 7: 'Firma'}
CAB_TIEMPO = ('Hora Límite', 'Hora', 'Día', 'Cadencia', 'Antelación', 'Cuándo')
RX_NUM = re.compile(r'^=COUNTIFS\(B5:B(\d+),"\?\*",F5:F\1,"✓"\)$')
RX_DEN = re.compile(r'^=COUNTIF\(B5:B(\d+),"\?\*"\)-COUNTIF\(F5:F\1,"N/A"\)$')
RX_PIE = re.compile(r'^— Kit de Tareas Recurrentes · .+ · AI Chef Pro · aichef\.pro$')
#: La fila 1 y la fila 2 llevan el nombre del kit; lo que se compara es la FORMA.
RX_R1 = re.compile(r'^(Checklist|Plantilla en Blanco)\b')
RX_R2 = re.compile(r'^(Fecha: ___/___/______|Semana del ___/___/______|'
                   r'Mes: ________________|Evento / temporada:|Año: ______)')


def _mapa(fila):
    return {c['c']: c for c in fila['cells']}


def _fill(cel):
    return (cel or {}).get('fill')


def comparar(path, titulo=None):
    fallos = []
    wb_ref = load_workbook(REF)
    ref = hoja(wb_ref[REF_HOJA])
    wb = load_workbook(path)
    nombre = titulo or wb.sheetnames[1]
    if nombre not in wb.sheetnames:
        return ['la hoja «%s» no existe en %s' % (nombre, path)]
    ws = wb[nombre]
    # BONUS-01 (briefing) y BONUS-02 (calendario) NO son del molde checklist: tienen el suyo
    # (contrato §5 y §6). Se dice en voz alta en vez de escupir 20 «fallos» que no lo son —y en
    # vez de saltárselo en silencio, que es cómo un gate se vuelve un falso verde.
    if ws.cell(row=4, column=2).value != 'Tarea':
        return None
    got = hoja(ws)

    # --- panel congelado --------------------------------------------------
    if got['freeze'] != ref['freeze']:
        fallos.append('freeze_panes %s ≠ %s (referencia)' % (got['freeze'], ref['freeze']))

    # --- anchos -----------------------------------------------------------
    for col in ('A', 'B', 'C', 'D', 'E', 'F', 'G'):
        a, b = got['widths'].get(col), ref['widths'].get(col)
        if a != b:
            fallos.append('ancho %s = %s ≠ %s (referencia)' % (col, a, b))

    # --- filas 1-4: texto, fill, bold --------------------------------------
    r1, r1r = _mapa(got['rows'][0]), _mapa(ref['rows'][0])
    a1 = r1.get('A')
    if not a1 or not isinstance(a1['v'], str) or not RX_R1.match(a1['v']):
        fallos.append('A1 = %r: no arranca por «Checklist» ni «Plantilla en Blanco»'
                      % (a1 or {}).get('v'))
    elif not (a1['bold'] and a1['sz'] == r1r['A']['sz'] and a1['color'] == r1r['A']['color']):
        fallos.append('A1: bold/size/color %s/%s/%s ≠ %s/%s/%s (referencia)'
                      % (a1['bold'], a1['sz'], a1['color'],
                         r1r['A']['bold'], r1r['A']['sz'], r1r['A']['color']))
    if 'A1:G1' not in got['merged']:
        fallos.append('falta el merge A1:G1')

    r2, r2r = _mapa(got['rows'][1]), _mapa(ref['rows'][1])
    a2 = r2.get('A')
    if not a2 or not isinstance(a2['v'], str) or not RX_R2.match(a2['v']):
        fallos.append('A2 = %r: no es ninguna de las filas 2 canónicas del molde'
                      % (a2 or {}).get('v'))
    elif a2['sz'] != r2r['A']['sz'] or a2['color'] != r2r['A']['color']:
        fallos.append('A2: size/color %s/%s ≠ %s/%s (referencia)'
                      % (a2['sz'], a2['color'], r2r['A']['sz'], r2r['A']['color']))
    if 'A2:G2' not in got['merged']:
        fallos.append('falta el merge A2:G2')

    if got['rows'][2]['cells']:
        fallos.append('la fila 3 debe ir VACÍA (la cabecera va en la 4)')

    r4, r4r = _mapa(got['rows'][3]), _mapa(ref['rows'][3])
    for i in range(1, 8):
        col = get_column_letter(i)
        cel, rcel = r4.get(col), r4r.get(col)
        if not cel:
            fallos.append('falta la cabecera %s4' % col)
            continue
        esperado = CAB_FIJA.get(i)
        if esperado and cel['v'] != esperado:
            fallos.append('%s4 = %r ≠ %r (rótulo del molde)' % (col, cel['v'], esperado))
        if i == 5 and cel['v'] not in CAB_TIEMPO:
            fallos.append('E4 = %r: no es un rótulo de tiempo que el motor conserve %s'
                          % (cel['v'], list(CAB_TIEMPO)))
        for k in ('fill', 'bold', 'color', 'sz'):
            if cel[k] != rcel[k]:
                fallos.append('%s4 %s = %r ≠ %r (referencia)' % (col, k, cel[k], rcel[k]))

    # --- validación de datos ----------------------------------------------
    if not got['dv']:
        fallos.append('la hoja no tiene ninguna validación de datos')
    else:
        for dv, rdv in zip(got['dv'], ref['dv']):
            if dv['type'] != rdv['type'] or dv['formula1'] != rdv['formula1']:
                fallos.append('DV %s/%r ≠ %s/%r (referencia)'
                              % (dv['type'], dv['formula1'], rdv['type'], rdv['formula1']))
            if dv['allow_blank'] != rdv['allow_blank']:
                fallos.append('DV allow_blank %s ≠ %s' % (dv['allow_blank'], rdv['allow_blank']))
        if len(got['dv']) != 1:
            fallos.append('%d validaciones de datos: el motor deja exactamente 1'
                          % len(got['dv']))

    # --- contador, filas libres, verificado y pie --------------------------
    contador = None
    for r in range(5, ws.max_row + 1):
        if ws.cell(row=r, column=1).value == 'Tareas completadas:':
            contador = r
            break
    if contador is None:
        fallos.append('no hay fila «Tareas completadas:»')
    else:
        num, den = ws.cell(row=contador, column=4).value, ws.cell(row=contador, column=6).value
        if not (isinstance(num, str) and RX_NUM.match(num)):
            fallos.append('numerador %r: no casa con COUNTIFS(B5:B<n>,"?*",F5:F<n>,"✓")' % num)
        if not (isinstance(den, str) and RX_DEN.match(den)):
            fallos.append('denominador %r: no casa con COUNTIF(B5:B<n>,"?*")-COUNTIF(F5:F<n>,"N/A")'
                          % den)
        if ws.cell(row=contador, column=5).value != 'de':
            fallos.append('E%d = %r ≠ «de»' % (contador, ws.cell(row=contador, column=5).value))
        n = int(RX_NUM.match(num).group(1)) if isinstance(num, str) and RX_NUM.match(num) else None
        if n is not None and n != contador - 2:
            fallos.append('el rango del contador acaba en %d y la fila de totales está en %d: '
                          'debe haber EXACTAMENTE una fila en blanco entre medias' % (n, contador))
        # 5 filas libres verdes SIN número justo antes
        primera = (n or contador - 2) - HOLGURA + 1
        for r in range(primera, (n or contador - 2) + 1):
            if ws.cell(row=r, column=1).value is not None:
                fallos.append('la fila libre %d lleva número en A (geometria la contaría '
                              'como tarea y cada pasada añadiría 5 más)' % r)
            for c in range(1, 8):
                cel = ws.cell(row=r, column=c)
                rgb = cel.fill.fgColor.rgb if cel.fill.fill_type == 'solid' else None
                if not (isinstance(rgb, str) and rgb.upper().endswith(VERDE)):
                    fallos.append('la fila libre %d no está verde en %s (%s)'
                                  % (r, get_column_letter(c), rgb))
                    break
        # «Verificado por / Firma»
        verif = contador + 2
        if ws.cell(row=verif, column=1).value != 'Verificado por:' \
                or ws.cell(row=verif, column=5).value != 'Firma:':
            fallos.append('falta la fila «Verificado por: / Firma:» en %d' % verif)

    pie = None
    for r in range(ws.max_row, 4, -1):
        v = ws.cell(row=r, column=1).value
        if isinstance(v, str) and v.startswith('—'):
            pie = v
            break
    if pie is None or not RX_PIE.match(pie):
        fallos.append('pie %r: no casa con «— Kit de Tareas Recurrentes · <kit> · '
                      'AI Chef Pro · aichef.pro»' % pie)

    if not got['protected']:
        fallos.append('la hoja NO está protegida (ws.protection.sheet)')
    if ws.protection.password is not None:
        fallos.append('la protección lleva contraseña; el molde va SIN contraseña')
    if not got['print_area']:
        fallos.append('print_area vacío')
    if not got['cf']:
        fallos.append('no hay formato condicional (fila verde al marcar ✓)')
    return fallos


def main():
    if len(sys.argv) < 2:
        raise SystemExit(__doc__)
    path = sys.argv[1]
    titulo = sys.argv[2] if len(sys.argv) > 2 else None
    fallos = comparar(path, titulo)
    ref = '%s!%s' % (os.path.basename(REF), REF_HOJA)
    if fallos is None:
        print('N/A  — %s: la hoja comparada no es del molde CHECKLIST (los dos BONUS '
              'tienen el suyo: briefing y calendario). Este gate no aplica.'
              % os.path.basename(path))
        return 3
    if fallos:
        print('FAIL — %s contra %s' % (os.path.basename(path), ref))
        for f in fallos:
            print('  · ' + f)
        return 1
    print('PASS — %s: mismo molde estructural que %s' % (os.path.basename(path), ref))
    return 0


if __name__ == '__main__':
    sys.exit(main())
