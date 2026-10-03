#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
aplicar_en.py — Restaurant Financial Plan Kit Pro (EN) · montaje de los 10 xlsx desde el Kit Plan Financiero para
Restaurantes v2.0 publicado.

SPEC que manda: `SPEC.md` (§2 mapas, §3 reglas US/UK, §4-§5, §7 pasos 3-5, D1-D24). Sesión Claude Code, 3-oct-2026.
Copia adaptada de `staff-kit/aplicar_en.py` (misma maquinaria: renombrado de hojas con reescritura de referencias,
literales en fórmulas/CF/DV, traducción POR APARICIÓN, formatos US, US Letter + pie, docProps, comillas rectas,
huella de idempotencia, `inject_cache` AL FINAL, `--dry-run`, `--idempotencia`) + `restaurar_graficos` de
`recipe-costing-kit/aplicar_en.py` (openpyxl reescribe el gráfico sin su <c:style>: se vuelve al XML del ES con las
referencias de pestaña y los títulos EN) + la caché de la TIR (pycel no evalúa `IRR`). Aquí no se redacta nada: los
textos salen de `textos_en/G*.json` (subagentes Anthropic, regla 1bis) y de `mapas.py`.

    python3 aplicar_en.py                      # escribe en astro-site/public/dl/restaurant-financial-plan-templates/
    python3 aplicar_en.py --dry-run [--salida <carpeta>] [--mes October] [--idempotencia] [--json informe.json]
    python3 aplicar_en.py --dry-run --parcial  # deja en español lo que aún no tenga traducción (solo ensayo)

Los ficheros ES de `dl/kit-plan-financiero/` se abren en SOLO LECTURA y nunca se escriben.

Orden por libro, siempre desde una copia limpia del ES publicado (→ idempotente):
   1. carga del xlsx ES y renombrado de hojas (mapas.HOJAS) reescribiendo las referencias de fórmulas y DV;
   2. fórmulas: el patrón de mapas.FORMULAS_EN celda a celda (D10: dotación = coste / vida útil, 48 celdas); el resto,
      HOJAS + CLAVES (un literal con letras sin clave aborta);
   3. textos POR APARICIÓN (textos_es.json `donde`): traducir → textos_en, mapa-clave → CLAVES, mapa-hoja → HOJAS;
      línea de versión, línea de marca y pie (D20, D18);
   4. DV: lista del B09 por CLAVES (D7), DV_PARTIDAS (D10: la DV 0-1 de C+H se parte; H = vida útil 0-50 años),
      referencias de hoja y mensajes por aparición (+ DV_MENSAJES_EN);
   5. CF: literales de la fórmula y atributo `text` por CLAVES;
   6. capas regeneradas: título interior (D5), POR_CELDA (D15), FIJOS (D8-D16, por texto), VALORES_EN (D8-D15),
      AJUSTE_TEXTO (ajuste + alto de fila) y ANCHOS (columnas de rótulos que el EN desborda);
   7. formatos (D17-D18): sin «€», fechas numFmtId 14, «p.p.» → 0.0 % (I1), « años» → « years», 04 H sin %
      (FORMATOS_POR_RANGO); US Letter en las 56 hojas; formatos propios sin uso renombrados;
   8. docProps D20, título interior comprobado, comillas rectas;
   9. cobertura: toda aparición «regenerar» del ES tiene que haberla escrito una capa;
  10. guardado con el nombre EN (mapas.FICHEROS) y restauración de los gráficos desde el XML del ES (D19);
  11. inject_cache AL FINAL sobre los 10, con `IRR` enseñada a pycel (misma Newton-Raphson que `cachear_irr` del ES:
      0 fallos de pycel), y `cachear_irr` del 07 que recalcula la TIR en Python desde los flujos cacheados y la
      compara con la caché (o la inyecta si faltara).

La protección de las 56 hojas (sin contraseña) viaja intacta con openpyxl.
Térmica: todo en SERIE; `istats` antes de cada libro y de cada inject_cache (espera si ≥ 62 °C hasta < 60 °C). En el
VPS (D24) no hay istats y la guarda no espera.
"""
import argparse
import copy
import datetime
import glob
import hashlib
import html
import json
import logging
import os
import re
import shutil
import subprocess
import sys
import tempfile
import time
import zipfile
from collections import OrderedDict
from xml.sax.saxutils import escape as xml_escape

AQUI = os.path.dirname(os.path.abspath(__file__))
SCRIPTS = os.path.dirname(AQUI)
REPO = os.path.dirname(os.path.dirname(SCRIPTS))
sys.path.insert(0, AQUI)
sys.path.insert(1, SCRIPTS)
sys.path.insert(2, os.path.join(SCRIPTS, 'kit-plan-financiero-v2_0'))

import openpyxl                                              # noqa: E402
from openpyxl.worksheet.cell_range import MultiCellRange     # noqa: E402
from openpyxl.worksheet.datavalidation import DataValidation  # noqa: E402

import mapas                                                 # noqa: E402
import extraer_textos                                        # noqa: E402

logging.disable(logging.CRITICAL)

ORIGEN = os.path.join(REPO, 'astro-site', 'public', 'dl', 'kit-plan-financiero')
DESTINO_REAL = os.path.join(REPO, 'astro-site', 'public', 'dl', mapas.SLUG)   # restaurant-financial-plan-templates
INJECT = os.path.join(SCRIPTS, 'inject_cache.py')

VERSION = '2.0'
SUBJECT = mapas.PRODUCTO + ' · v2.0'                                  # D20
DESCRIPTION = mapas.URL_PRODUCTO
KEYWORDS = ('restaurant p&l template, restaurant financial projections, restaurant budget template, break-even, '
            'cash flow, AI Chef Pro')
CATEGORY = 'AI Chef Pro · Digital products'
CREATOR = 'AI Chef Pro'
PIE_EN = mapas.FOOTER_EN                                              # «AI Chef Pro · aichef.pro · Page &P of &N»
MESES = ['January', 'February', 'March', 'April', 'May', 'June', 'July', 'August',
         'September', 'October', 'November', 'December']
VERDE = 'E8F5E9'
FMT_FECHA = 'mm-dd-yy'                     # numFmtId 14 (fecha corta del sistema: m/d/yyyy en US)
NUMFMT = {FMT_FECHA: 14}
FORMATOS = {k: (FMT_FECHA if v == 'NUMFMT14' else v) for k, v in mapas.FORMATOS_EN.items()}
INVARIABLES = set(extraer_textos.INVARIABLES)
RX_VERSION_ES = extraer_textos.RX_VERSION
HOJA_INSTR_EN = mapas.HOJAS['Instrucciones']
CLAVES = mapas.CLAVES
CAMPO_DV = {'dv-error-titulo': 'errorTitle', 'dv-error': 'error', 'dv-prompt-titulo': 'promptTitle', 'dv-prompt': 'prompt'}

# D12: la TIR del 07 (pycel no implementa IRR) — hoja EN, celda, rango de flujos y tasa del VAN
CELDA_IRR = (mapas.HOJAS['Proyecciones'], 'B24')
RANGO_FLUJOS = (mapas.HOJAS['Proyecciones'], 'B22:G22')
CELDA_TASA = (mapas.HOJAS['Proyecciones'], 'C27')
FICHERO_07 = mapas.FICHEROS['07-informe-viabilidad-bancos.xlsx']


class Aborta(Exception):
    pass


# ==========================================================================
# Térmica (regla del Mac: se apaga por encima de ~65 °C; en el VPS no hay istats)
# ==========================================================================
def termica(umbral=62, reanuda=60):
    def leer():
        try:
            out = subprocess.run(['istats', 'cpu', 'temp', '--value-only'], stdout=subprocess.PIPE,
                                 stderr=subprocess.PIPE, universal_newlines=True, timeout=20).stdout
            return float(re.findall(r'[\d.]+', out)[0])
        except Exception:
            return None
    t = leer()
    if t is None or t < umbral:
        return t
    while t is not None and t >= reanuda:
        print('  [térmica] CPU %.1f °C: espero 30 s' % t, flush=True)
        time.sleep(30)
        t = leer()
    return t


# ==========================================================================
# TIR (D12): Newton-Raphson del ES (`kit-plan-financiero-v2_0/motor.py`), enseñada a pycel e inyectada
# ==========================================================================
def _motor():
    import motor                                              # kit-plan-financiero-v2_0 (solo funciones puras)
    return motor


def tir(flujos):
    """TIR por Newton-Raphson (motor.tir_newton del ES); None si no hay cambio de signo (Excel: #NUM!)."""
    return _motor().tir_newton(flujos)


_PYCEL_PARCHEADO = []


def parchear_pycel():
    """pycel 1.0b30 no implementa IRR («Function IRR is not implemented») y la celda se quedaría sin caché en el
    documento que se entrega al lender. Se le enseña la MISMA Newton-Raphson del ES: valores no numéricos fuera
    (como Excel con referencias), #NUM! sin cambio de signo (el IFERROR de la fórmula devuelve entonces «—»)."""
    if _PYCEL_PARCHEADO:
        return
    import pycel.excellib as XL
    from pycel.excelutil import flatten, is_number, NUM_ERROR

    def irr(values, guess=None):                              # noqa: ARG001 (guess: Newton parte de 0.1 como el ES)
        flujos = [float(x) for x in flatten(values) if is_number(x) and not isinstance(x, bool)]
        r = tir(flujos)
        return NUM_ERROR if r is None else r
    XL.irr = irr
    _PYCEL_PARCHEADO.append(True)


def inyectar_cache(path):
    """inject_cache.inject en el MISMO proceso con IRR parcheada → (fórmulas, inyectadas, fallos_pycel)."""
    parchear_pycel()
    import inject_cache as IC
    return IC.inject(path)


def leer_flujos(path):
    wv = openpyxl.load_workbook(path, data_only=True)
    ws = wv[RANGO_FLUJOS[0]]
    fl = []
    for c in ws[RANGO_FLUJOS[1]][0]:
        fl.append(float(c.value) if isinstance(c.value, (int, float)) and not isinstance(c.value, bool) else 0.0)
    return fl, ws[CELDA_IRR[1]].value


def cachear_irr(path):
    """Tras inject_cache: TIR en Python desde los flujos CACHEADOS; compara con la caché de B24 (o la inyecta)."""
    flujos, cache = leer_flujos(path)
    t = tir(flujos)
    formula = openpyxl.load_workbook(path)[CELDA_IRR[0]][CELDA_IRR[1]].value
    m = re.search(r',\s*"([^"]*)"\s*\)\s*$', formula or '')
    alterna = m.group(1) if m else '—'
    esperado = t if t is not None else alterna
    out = OrderedDict([('flujos', flujos), ('tir_python', t), ('cache_pycel', cache), ('inyectada', False)])
    if cache is None:
        out['inyectada'] = _motor().inyectar_valor(path, CELDA_IRR[0], CELDA_IRR[1], esperado)
        if not out['inyectada']:
            raise Aborta('%s: no se pudo inyectar la TIR' % os.path.basename(path))
    elif t is None:
        if cache != alterna:
            raise Aborta('%s: sin cambio de signo la TIR cacheada debe ser %r y es %r' % (path, alterna, cache))
    elif not isinstance(cache, (int, float)) or abs(cache - t) > 1e-9:
        raise Aborta('%s: TIR cacheada %r ≠ Newton %r' % (path, cache, t))
    out['valor'] = esperado
    return out


# ==========================================================================
# Datos
# ==========================================================================
def cargar_datos(parcial=False):
    te = json.load(open(os.path.join(AQUI, 'textos_es.json'), encoding='utf-8'))
    censo = json.load(open(os.path.join(AQUI, 'censo_es.json'), encoding='utf-8'))
    tx = OrderedDict()
    for p in sorted(glob.glob(os.path.join(AQUI, 'textos_en', 'G[0-9].json'))):
        for k, v in json.load(open(p, encoding='utf-8')).items():
            if k in tx and tx[k] != v:
                raise Aborta('traducción contradictoria para %s en %s' % (k, os.path.basename(p)))
            tx[k] = v
    por_id = {c['id']: c for c in te['cadenas']}
    ajenas = sorted(set(tx) - set(por_id))
    if ajenas:
        raise Aborta('textos_en trae ids que no están en textos_es.json: %r' % ajenas[:5])
    gm = sorted(k for k in tx if por_id[k]['grupo'] == 'GM')
    if gm:
        raise Aborta('textos_en traduce cadenas GM (las escribe mapas.py): %r' % gm[:5])
    faltan = [c['id'] for c in te['cadenas'] if c['grupo'] != 'GM' and c['id'] not in tx]
    if faltan and not parcial:
        raise Aborta('%d cadenas G1-G2 sin traducción (p. ej. %r)' % (len(faltan), faltan[:5]))
    celdas, dvmsg, graf = {}, {}, {}
    for c in te['cadenas']:
        for d in c['donde']:
            if d['t'] == 'celda':
                celdas[(d['f'], d['h'], d['c'])] = (c['id'], d['tr'], c['es'])
            elif d['t'] in CAMPO_DV:
                dvmsg[(d['f'], d['h'], d['c'], CAMPO_DV[d['t']])] = (c['id'], d['tr'], c['es'])
            elif d['t'] in ('grafico-titulo', 'grafico-eje'):
                graf.setdefault(d['f'], {})[c['es']] = (c['id'], d['tr'], c['es'])
    return {'te': te, 'censo': censo, 'tx': tx, 'celdas': celdas, 'dvmsg': dvmsg, 'graf': graf, 'faltan': faltan,
            'parcial': parcial}


def mes_por_defecto():
    return MESES[datetime.date.today().month - 1]


def linea_version(mes):
    return 'Version %s · %s %d · %s · info@aichef.pro' % (VERSION, mes, datetime.date.today().year, DESCRIPTION)


def titulo_docprops(fname_es):
    """D5: «NN · título · Restaurant Financial Plan Kit Pro» (como el ES: título interior + producto)."""
    return mapas.titulo_interior(fname_es) + ' · ' + mapas.PRODUCTO


# ==========================================================================
# Fórmulas: hojas y literales
# ==========================================================================
RX_STR = re.compile(r'"(?:[^"]|"")*"')
RX_REF = re.compile(r"(?:'((?:[^']|'')+)'|([^\W\d][\w.]*))!")


def ref_hoja(nombre):
    """Nombre de hoja tal como va en una fórmula (entre comillas si hace falta)."""
    if re.fullmatch(r'[A-Za-z_][A-Za-z0-9_.]*', nombre):
        return nombre
    return "'" + nombre.replace("'", "''") + "'"


def clave_en(lit, donde):
    """Literal ES → EN por CLAVES: estricto con los que llevan letras o «€»; los demás no cambian."""
    if not mapas.necesita_clave(lit):
        return lit
    if lit not in CLAVES:
        raise Aborta('literal sin clave EN %r en %r' % (lit, donde[:90]))
    return CLAVES[lit]


class Transformador:
    """Reescribe una fórmula ES a EN: referencias de hoja (fuera de literales) y literales (dentro)."""

    def __init__(self, mapa_hojas):
        self.mapa = mapa_hojas
        self.en = set(mapa_hojas.values())
        self.hojas_citadas = 0
        self.literales = 0

    def refs(self, s):
        def rep(m):
            nombre = (m.group(1) if m.group(1) is not None else m.group(2)).replace("''", "'")
            if nombre in self.mapa:
                self.hojas_citadas += 1
                return ref_hoja(self.mapa[nombre]) + '!'
            if nombre in self.en:
                return m.group(0)
            raise Aborta('hoja desconocida en fórmula: %r en %r' % (nombre, s[:80]))
        return RX_REF.sub(rep, s)

    def formula(self, f):
        out, pos = [], 0
        for m in RX_STR.finditer(f):
            out.append(self.refs(f[pos:m.start()]))
            lit = m.group(0)[1:-1].replace('""', '"')
            nuevo = clave_en(lit, f)
            if nuevo != lit:
                self.literales += 1
            out.append('"' + nuevo.replace('"', '""') + '"')
            pos = m.end()
        out.append(self.refs(f[pos:]))
        return ''.join(out)

    def celda(self, v):
        return '=' + self.formula(v[1:]) if v.startswith('=') else self.formula(v)


def formula_en(v, fila, T):
    """Fórmula EN de una celda: patrón de FORMULAS_EN (D10) o HOJAS + CLAVES. Devuelve (nueva, es_patron)."""
    pat = mapas.patron(v)
    if pat in mapas.FORMULAS_EN:
        if pat.replace('{r}', str(fila)) != v:
            raise Aborta('patrón FORMULAS_EN que no es de una sola fila: %r' % v[:90])
        return mapas.FORMULAS_EN[pat].replace('{r}', str(fila)), True
    return T.celda(v), False


def rango(ref):
    c1, r1, c2, r2 = openpyxl.utils.cell.range_boundaries(ref)
    for r in range(r1, r2 + 1):
        for c in range(c1, c2 + 1):
            yield '%s%d' % (openpyxl.utils.get_column_letter(c), r)


def rango_lista(sqref):
    out = []
    for rg in str(sqref).split():
        out.extend(rango(rg))
    return out


def renombrar(wb):
    mapa = OrderedDict((ws.title, mapas.HOJAS[ws.title]) for ws in wb.worksheets)
    for ws in wb.worksheets:
        ws.title = mapa[ws.title]
    if wb.sheetnames != list(mapa.values()):
        raise Aborta('renombrado: %r ≠ %r' % (wb.sheetnames, list(mapa.values())))
    return mapa


def reescribir_formulas(wb, T, rep):
    n = pat = 0
    for ws in wb.worksheets:
        for c in list(ws._cells.values()):
            if c.data_type == 'f' and isinstance(c.value, str):
                nuevo, es_pat = formula_en(c.value, c.row, T)
                pat += es_pat
                if nuevo != c.value:
                    c.value = nuevo
                    n += 1
    rep['formulas_reescritas'] = n
    rep['formulas_patron_en'] = pat


def reescribir_titulos_impresion(wb, mapa, rep):
    """Los títulos de impresión viajan como nombres definidos locales; openpyxl los regenera con el título EN de
    la hoja. Se comprueba (no se parchea): ninguno puede seguir citando una pestaña ES."""
    n = 0
    for ws in wb.worksheets:
        for attr in ('print_title_rows', 'print_title_cols'):
            v = getattr(ws, attr)
            if v:
                n += 1
                if any(es in str(v) for es in mapa if es not in mapa.values()):
                    raise Aborta('%s: título de impresión con pestaña ES %r' % (ws.title, v))
    rep['titulos_impresion'] = n


# ==========================================================================
# Textos por aparición
# ==========================================================================
class Libro:
    """Estado de un libro en construcción: celdas escritas por capas (cobertura de «regenerar»)."""

    def __init__(self, wb, corto, mapa, datos, rep):
        self.wb, self.corto, self.mapa, self.datos, self.rep = wb, corto, mapa, datos, rep
        self.inv = {en: es for es, en in mapa.items()}
        self.escritas = set()                 # (hoja_es, celda)
        self.faltan = rep.setdefault('faltan', [])

    def ws(self, hoja_es):
        return self.wb[self.mapa[hoja_es]]

    def poner(self, hoja_es, coord, valor, es_esperado=None):
        c = self.ws(hoja_es)[coord]
        if c.data_type == 'f' or (isinstance(c.value, str) and c.value.startswith('=')):
            raise Aborta('%s %s!%s es una fórmula (%r): una capa no la pisa' % (self.corto, hoja_es, coord, c.value))
        if es_esperado is not None and not es_esperado(c.value):
            raise Aborta('%s %s!%s: el ES trae %r, no lo esperado' % (self.corto, hoja_es, coord, c.value))
        c.value = valor
        self.escritas.add((hoja_es, coord))
        return c

    def en_de(self, ident, tr, es, lugar):
        if tr in ('traducir', 'localizar'):
            en = self.datos['tx'].get(ident)
            if en is None:
                self.faltan.append('%s %s %r' % (lugar, ident, es[:60]))
                return es
            return en
        if tr == 'mapa-clave':
            return CLAVES[es]
        if tr == 'mapa-hoja':
            return mapas.HOJAS[es]
        return None                                              # regenerar: lo escribe una capa


def es_texto(v):
    return isinstance(v, str) and bool(v)


def traducir_textos(L, mes):
    n = 0
    regenerar = []
    for ws in L.wb.worksheets:
        h_es = L.inv[ws.title]
        for c in list(ws._cells.values()):
            v = c.value
            if not isinstance(v, str) or c.data_type == 'f':
                continue
            if RX_VERSION_ES.match(v):
                L.poner(h_es, c.coordinate, linea_version(mes))
                n += 1
                continue
            if v == mapas.MARCA_ES:
                L.poner(h_es, c.coordinate, mapas.MARCA_EN)
                n += 1
                continue
            k = (L.corto, h_es, c.coordinate)
            if k not in L.datos['celdas']:
                if (any(ch.isalpha() for ch in v) or '€' in v) and v not in INVARIABLES:
                    raise Aborta('%s %s!%s: texto %r fuera de textos_es.json' % (L.corto, h_es, c.coordinate, v[:60]))
                continue                                         # sin letras / invariable (▸, —, 25-28%…)
            ident, tr, es = L.datos['celdas'][k]
            if es != v:
                raise Aborta('%s %s!%s: textos_es.json dice %r y el ES trae %r' % (L.corto, h_es, c.coordinate, es[:40], v[:40]))
            en = L.en_de(ident, tr, es, '%s %s!%s' % (L.corto, ws.title, c.coordinate))
            if en is None:
                regenerar.append((h_es, c.coordinate))
                continue
            if en != v:
                c.value = en
                n += 1
        for parte in ('oddHeader', 'oddFooter', 'evenHeader', 'evenFooter', 'firstHeader', 'firstFooter'):
            hf = getattr(ws, parte)
            for pos in ('left', 'center', 'right'):
                t = getattr(hf, pos).text
                if not t:
                    continue
                if t != mapas.FOOTER_ES:
                    raise Aborta('%s %s %s.%s inesperado: %r' % (L.corto, ws.title, parte, pos, t))
                getattr(hf, pos).text = PIE_EN
    L.rep['textos_traducidos'] = n
    return regenerar


# ==========================================================================
# DV y CF
# ==========================================================================
def dv_censo(L, h_es, sq):
    for d in L.datos['censo']['libros'][L.rep['es']]['hojas_detalle'][h_es]['dv']:
        if d['sqref'] == sq:
            return d
    raise Aborta('%s %s: DV %s no está en el censo' % (L.corto, h_es, sq))


def dv_lista_en(items, donde):
    out = []
    for x in items:
        if x == '' or not mapas.necesita_clave(x):
            out.append(x)
            continue
        if x not in CLAVES:
            raise Aborta('%s: ítem de DV sin clave EN %r' % (donde, x))
        out.append(CLAVES[x])
    return mapas.dv_lista(out)


def dv_partida_nueva(dv_es, partida):
    """D10: DV nueva de la columna H (vida útil en años) con los flags de la DV ES de la que sale."""
    _conserva, sq_nueva, tipo, op, f1, f2, msg, titulo = partida
    nueva = DataValidation(type=tipo, operator=op, formula1=f1, formula2=f2, allow_blank=dv_es.allow_blank,
                           showDropDown=dv_es.showDropDown, showErrorMessage=dv_es.showErrorMessage,
                           showInputMessage=dv_es.showInputMessage, errorStyle=dv_es.errorStyle,
                           error=msg, errorTitle=titulo)
    nueva.sqref = MultiCellRange(sq_nueva)
    return nueva


def reescribir_dv_cf(L, T):
    n = OrderedDict([('dv_listas', 0), ('dv_personalizadas', 0), ('dv_ref', 0), ('dv_partidas', 0),
                     ('dv_mensajes', 0), ('dv_mensajes_fijados', 0), ('cf_reglas', 0), ('cf_text', 0)])
    usados = set()
    for ws in L.wb.worksheets:
        h_es = L.inv[ws.title]
        nuevas = []
        for dv in list(ws.data_validations.dataValidation):
            sq = str(dv.sqref)
            f1 = dv.formula1 or ''
            dv_censo(L, h_es, sq)
            if dv.type == 'list' and f1.startswith('"'):
                dv.formula1 = dv_lista_en(f1.strip('"').split(','), '%s %s DV %s' % (L.corto, h_es, sq))
                if len(dv.formula1) > 257:
                    raise Aborta('%s %s: DV %s > 255 caracteres' % (L.corto, h_es, sq))
                n['dv_listas'] += 1
            elif dv.type == 'custom' and f1:
                dv.formula1 = T.formula(f1)
                n['dv_personalizadas'] += 1
            elif f1:
                dv.formula1 = T.refs(f1)
                if dv.formula2:
                    dv.formula2 = T.refs(dv.formula2)
                n['dv_ref'] += 1
            for attr in ('error', 'errorTitle', 'prompt', 'promptTitle'):
                v = getattr(dv, attr)
                if not v:
                    continue
                k = (L.corto, h_es, sq, attr)
                if k not in L.datos['dvmsg']:
                    raise Aborta('%s %s DV %s %s fuera de textos_es.json: %r' % (L.corto, h_es, sq, attr, v[:50]))
                ident, tr, es = L.datos['dvmsg'][k]
                if es != v:
                    raise Aborta('%s %s DV %s %s: textos_es.json dice %r' % (L.corto, h_es, sq, attr, es[:40]))
                en = L.en_de(ident, tr, es, '%s DV %s %s' % (ws.title, sq, attr))
                if en is None:
                    raise Aborta('%s %s DV %s %s marcado «regenerar»' % (L.corto, h_es, sq, attr))
                fijo = mapas.DV_MENSAJES_EN.get((L.corto, h_es, sq, attr))
                if fijo is not None:                       # revisión final: arreglo solo para ESTA aparición
                    en = fijo
                    usados.add((L.corto, h_es, sq, attr))
                    n['dv_mensajes_fijados'] += 1
                setattr(dv, attr, en)
                n['dv_mensajes'] += 1
            partida = mapas.DV_PARTIDAS.get((L.corto, h_es, sq))
            if partida:                                    # D10: C conserva la DV 0-1; H estrena la de vida útil
                dv.sqref = MultiCellRange(partida[0])
                nuevas.append(dv_partida_nueva(dv, partida))
                n['dv_partidas'] += 1
        for nueva in nuevas:
            ws.add_data_validation(nueva)
        for cf in ws.conditional_formatting:
            for regla in cf.rules:
                if regla.formula:
                    regla.formula = [T.formula(x) for x in regla.formula]
                    n['cf_reglas'] += 1
                if regla.text:
                    regla.text = clave_en(regla.text, 'CF %s %s' % (ws.title, cf.sqref))
                    n['cf_text'] += 1
    sin_usar = [k for k in mapas.DV_MENSAJES_EN if k[0] == L.corto and k not in usados]
    if sin_usar:
        raise Aborta('%s: DV_MENSAJES_EN sin DV donde aplicarse: %r' % (L.corto, sin_usar))
    esperadas = sum(1 for k in mapas.DV_PARTIDAS if k[0] == L.corto)
    if n['dv_partidas'] != esperadas:
        raise Aborta('%s: %d DV partidas ≠ %d de DV_PARTIDAS' % (L.corto, n['dv_partidas'], esperadas))
    L.rep.update(n)


# ==========================================================================
# Capas regeneradas (mapas.py)
# ==========================================================================
def capa_titulo(L, fname_es):
    L.poner('Instrucciones', 'B2', mapas.titulo_interior(fname_es), es_texto)


def capa_por_celda(L):
    n = 0
    for (c, h, coord), v in mapas.POR_CELDA.items():
        if c == L.corto:
            L.poner(h, coord, v, es_texto)
            n += 1
    L.rep['por_celda'] = n


def capa_fijos(L, regenerar):
    """FIJOS (D8-D16): cadenas con decisión de mercado, por TEXTO, en toda aparición «regenerar» aún sin escribir."""
    n = 0
    for h_es, coord in regenerar:
        if (h_es, coord) in L.escritas:
            continue
        v = L.ws(h_es)[coord].value
        if v in mapas.FIJOS:
            L.poner(h_es, coord, mapas.FIJOS[v], es_texto)
            n += 1
    L.rep['fijos'] = n


def capa_valores(L):
    n = 0
    num = lambda v: isinstance(v, (int, float, datetime.date)) and not isinstance(v, bool)    # noqa: E731
    for (c, h, ref), v in mapas.VALORES_EN.items():
        if c != L.corto:
            continue
        celdas = list(rango(ref))
        if len(celdas) == 1:
            L.poner(h, ref, v, num)
            n += 1
            continue
        ws = L.ws(h)
        con = [x for x in celdas if ws[x].value is not None]
        if not con:
            raise Aborta('%s %s!%s: rango sin valores en el ES' % (c, h, ref))
        for coord in con:
            L.poner(h, coord, v, num)
            n += 1
    L.rep['valores'] = n


def alto_necesario(ws, fila):
    """Alto (pt) que pide el texto ajustado más largo de la fila. Estimación conservadora: un carácter por unidad
    de ancho de columna a 11 pt (el texto real es más estrecho que el «0» de referencia): sobra aire, no se corta."""
    alto = 0
    for cel in ws[fila]:
        v = cel.value
        if not isinstance(v, str) or cel.data_type == 'f' or not (cel.alignment and cel.alignment.wrap_text):
            continue
        if any(cel.coordinate in m for m in ws.merged_cells.ranges):
            continue
        ancho = ws.column_dimensions[cel.column_letter].width or 8.43
        pt = (cel.font.sz if cel.font is not None and cel.font.sz else 11)
        cpl = max(1, int(ancho * 11.0 / pt))
        lineas = sum(max(1, -(-len(seg) // cpl)) for seg in v.split('\n'))
        alto = max(alto, lineas * pt * 1.36)
    return round(alto * 4) / 4.0


def capa_ajuste(L):
    """Ajuste de texto en los rótulos EN que se cortaban (mapas.AJUSTE_TEXTO) y alto de fila."""
    n = 0
    for (c, h, rg) in mapas.AJUSTE_TEXTO:
        if c != L.corto:
            continue
        ws = L.ws(h)
        filas = set()
        for coord in rango(rg):
            cel = ws[coord]
            if cel.value is None:
                continue
            if not isinstance(cel.value, str) or cel.data_type == 'f':
                raise Aborta('%s %s!%s: AJUSTE_TEXTO sobre algo que no es un rótulo: %r' % (c, h, coord, cel.value))
            al = copy.copy(cel.alignment)
            al.wrap_text = True
            cel.alignment = al
            filas.add(cel.row)
            n += 1
        for fila in sorted(filas):
            alto = alto_necesario(ws, fila)
            actual = ws.row_dimensions[fila].height
            if alto > (actual or 15):
                ws.row_dimensions[fila].height = alto
    L.rep['ajuste_texto'] = n


def capa_anchos(L):
    """Columnas de rótulos que el EN desborda (mapas.ANCHOS): nunca estrecha una columna."""
    n = 0
    for (c, h), cols in mapas.ANCHOS.items():
        if c != L.corto:
            continue
        ws = L.ws(h)
        for col, w in cols.items():
            actual = ws.column_dimensions[col].width
            if actual is not None and actual > w:
                raise Aborta('%s %s: ANCHOS estrecharía la columna %s (%s → %s)' % (c, h, col, actual, w))
            ws.column_dimensions[col].width = w
            n += 1
    L.rep['anchos'] = n


# ==========================================================================
# Formatos y papel (D17, D18)
# ==========================================================================
RX_FMT_ES = re.compile(r'€|dd/mm|años|p\.p\.', re.I)


def formatos_y_papel(L):
    wb = L.wb
    n = OrderedDict([('fmt_' + k, 0) for k in FORMATOS])
    n['fmt_por_rango'] = 0
    for ws in wb.worksheets:
        for c in list(ws._cells.values()):
            f = c.number_format
            if f in FORMATOS:
                c.number_format = FORMATOS[f]
                n['fmt_' + f] += 1
        ws.page_setup.paperSize = ws.PAPERSIZE_LETTER
    for (c, h, rg), fmt in mapas.FORMATOS_POR_RANGO.items():
        if c != L.corto:
            continue
        ws = L.ws(h)
        for coord in rango(rg):
            ws[coord].number_format = fmt
            n['fmt_por_rango'] += 1
    # formatos propios que ya no usa ninguna celda (los ES) → variantes US sin €, dd/mm, «años» ni «p.p.»
    nf = wb._number_formats
    usados = set(nf)
    for i, f in enumerate(list(nf)):
        if RX_FMT_ES.search(f):
            nuevo = FORMATOS.get(f, '#,##0.00')
            if nuevo == FMT_FECHA:
                nuevo = 'm/d/yyyy'
            while nuevo in usados:
                nuevo += ';@'
            usados.add(nuevo)
            list.__setitem__(nf, i, nuevo)
    if len(set(nf)) != len(nf):
        raise Aborta('formatos propios repetidos: %r' % list(nf))
    nf._dict = {v: i for i, v in enumerate(nf)}
    nf.clean = True
    L.rep.update(n)


COMILLAS = {'\u2019': "'", '\u2018': "'", '\u201c': '"', '\u201d': '"'}


def norm_comillas(v):
    for a, b in COMILLAS.items():
        v = v.replace(a, b)
    return v


def normalizar_comillas(L):
    n = 0
    for ws in L.wb.worksheets:
        for c in list(ws._cells.values()):
            if isinstance(c.value, str) and c.data_type != 'f' and any(a in c.value for a in COMILLAS):
                c.value = norm_comillas(c.value)
                n += 1
        for dv in ws.data_validations.dataValidation:
            for attr in ('error', 'errorTitle', 'prompt', 'promptTitle'):
                v = getattr(dv, attr)
                if v and any(a in v for a in COMILLAS):
                    setattr(dv, attr, norm_comillas(v))
                    n += 1
    L.rep['comillas_normalizadas'] = n


def metadatos(L, fname_es):
    p = L.wb.properties
    p.title = titulo_docprops(fname_es)                # D5
    p.subject = SUBJECT
    p.keywords = KEYWORDS
    p.description = DESCRIPTION
    p.category = CATEGORY
    p.creator = CREATOR
    p.lastModifiedBy = CREATOR
    L.rep['docprops'] = OrderedDict((k, getattr(p, k)) for k in
                                    ('title', 'subject', 'creator', 'keywords', 'description', 'category'))


def comprobar_titulo(L, fname_es):
    """D5: el título interior (Instructions!B2) es EXACTAMENTE mapas.titulo_interior. Nunca se parchea."""
    esperado = mapas.titulo_interior(fname_es)
    v = L.wb[HOJA_INSTR_EN]['B2'].value
    if v != esperado:
        raise Aborta('%s: Instructions!B2 %r ≠ D5 %r' % (fname_es, v, esperado))
    L.rep['titulo'] = v


# ==========================================================================
# Gráficos (D19): XML del ES con referencias de pestaña EN y títulos EN
# ==========================================================================
RX_F = re.compile(r'(<(?:c:)?f>)([^<]*)(</(?:c:)?f>)')
RX_AT = re.compile(r'(<a:t>)([^<]*)(</a:t>)')
RX_FORMATCODE = re.compile(r'(formatCode=")([^"]*)(")')


def texto_grafico(datos, corto, t, faltan):
    if t in mapas.FIJOS:
        return mapas.FIJOS[t]
    if not any(ch.isalpha() for ch in t):
        return t                                                       # «%» (eje del 06)
    ent = datos['graf'].get(corto, {}).get(t)
    if ent is None:
        raise Aborta('%s: texto de gráfico %r fuera de textos_es.json' % (corto, t))
    en = datos['tx'].get(ent[0])
    if en is None:
        faltan.append('%s gráfico %s %r' % (corto, ent[0], t))
        return t
    return norm_comillas(en)


def restaurar_graficos(path_es, path_en, mapa, T, datos, corto, rep):
    """openpyxl reescribe el gráfico sin su <c:style>: se vuelve al XML original del ES con las referencias de
    pestaña EN (fórmulas de serie y de categoría), los títulos EN y formatos sin €. El ancla la conserva openpyxl."""
    g_es = extraer_textos.graficos(path_es)
    if not g_es:
        rep['graficos'] = []
        return
    g_en = extraer_textos.graficos(path_en)
    if len(g_en) != len(g_es):
        raise Aborta('%s: %d gráficos en EN y %d en ES' % (path_en, len(g_en), len(g_es)))
    parte_en = {}
    for g in g_en:
        if g['hoja'] in parte_en:
            raise Aborta('%s: dos gráficos en la hoja %s (el mapeo es por hoja)' % (path_en, g['hoja']))
        parte_en[g['hoja']] = g['parte']
    with zipfile.ZipFile(path_es) as z:
        nuevos = {}
        for g in g_es:
            x = z.read(g['parte']).decode('utf-8')
            x = RX_F.sub(lambda m: m.group(1) + xml_escape(T.formula(html.unescape(m.group(2)))) + m.group(3), x)
            x = RX_AT.sub(lambda m: m.group(1) + xml_escape(texto_grafico(datos, corto, html.unescape(m.group(2)),
                                                                          rep['faltan'])) + m.group(3), x)
            x = RX_FORMATCODE.sub(lambda m: m.group(1) + xml_escape(FORMATOS.get(html.unescape(m.group(2)),
                                                                                 html.unescape(m.group(2))))
                                  + m.group(3), x)
            destino = parte_en.get(mapa[g['hoja']])
            if not destino:
                raise Aborta('gráfico de «%s» sin parte EN' % g['hoja'])
            nuevos[destino] = x.encode('utf-8')
    with zipfile.ZipFile(path_en) as z:
        infos = z.infolist()
        partes = {i.filename: z.read(i.filename) for i in infos}
    partes.update(nuevos)
    tmp = path_en + '.tmp'
    with zipfile.ZipFile(tmp, 'w', zipfile.ZIP_DEFLATED) as zout:
        for i in infos:
            zout.writestr(i, partes[i.filename])
    os.replace(tmp, path_en)
    rep['graficos'] = sorted(nuevos)


# ==========================================================================
# Un libro
# ==========================================================================
def construir_libro(fname_es, carpeta, datos, mes):
    fname_en = mapas.FICHEROS[fname_es]
    corto = extraer_textos.corto(fname_es)
    src = os.path.join(ORIGEN, fname_es)
    dst = os.path.join(carpeta, fname_en)
    rep = OrderedDict([('es', fname_es), ('en', fname_en)])
    wb = openpyxl.load_workbook(src)
    mapa = renombrar(wb)                                                       # 1
    L = Libro(wb, corto, mapa, datos, rep)
    T = Transformador(mapa)
    reescribir_formulas(wb, T, rep)                                            # 2
    regenerar = traducir_textos(L, mes)                                        # 3
    reescribir_dv_cf(L, T)                                                     # 4-5
    capa_titulo(L, fname_es)                                                   # 6
    capa_por_celda(L)
    capa_fijos(L, regenerar)
    capa_valores(L)
    capa_ajuste(L)
    capa_anchos(L)
    rep['refs_hoja'] = T.hojas_citadas
    rep['literales'] = T.literales
    formatos_y_papel(L)                                                        # 7
    metadatos(L, fname_es)                                                     # 8
    comprobar_titulo(L, fname_es)
    normalizar_comillas(L)
    reescribir_titulos_impresion(wb, mapa, rep)
    sin_capa = [x for x in regenerar if x not in L.escritas]                   # 9
    if sin_capa:
        raise Aborta('%s: apariciones «regenerar» que ninguna capa escribió: %r' % (fname_es, sin_capa[:8]))
    rep['regeneradas'] = len(regenerar)
    wb.save(dst)                                                               # 10
    restaurar_graficos(src, dst, mapa, T, datos, corto, rep)
    if rep['faltan'] and not datos['parcial']:
        raise Aborta('%s: %d textos sin traducción:\n  %s' % (fname_es, len(rep['faltan']), '\n  '.join(rep['faltan'][:20])))
    return rep


def construir(carpeta, datos, mes, log=print):
    os.makedirs(carpeta, exist_ok=True)
    informe = []
    for fname_es in mapas.LIBROS_ES:
        termica()
        rep = construir_libro(fname_es, carpeta, datos, mes)
        informe.append(rep)
        log('  %-50s fórmulas %4d (EN %2d) · refs %4d · lit %3d · textos %4d · regen %3d · fijos %3d · dv÷ %d · '
            'gráf %d · faltan %d' % (rep['en'], rep['formulas_reescritas'], rep['formulas_patron_en'], rep['refs_hoja'],
                                     rep['literales'], rep['textos_traducidos'], rep['regeneradas'], rep['fijos'],
                                     rep['dv_partidas'], len(rep['graficos']), len(rep['faltan'])))
    return informe


# ==========================================================================
# Idempotencia (huella comparable, sin fechas de guardado)
# ==========================================================================
def digest(path):
    wb = openpyxl.load_workbook(path)
    p = wb.properties
    fuera = {'_props': (p.title, p.subject, p.keywords, p.description, p.category, p.creator)}
    with zipfile.ZipFile(path) as z:
        fuera['_graficos'] = sorted((n, hashlib.sha1(z.read(n)).hexdigest()) for n in z.namelist()
                                    if n.startswith('xl/charts/') or n.startswith('xl/drawings/'))
    for ws in wb.worksheets:
        celdas = {}
        for c in ws._cells.values():
            if c.value is None and not (c.fill is not None and c.fill.fill_type == 'solid'):
                continue
            relleno = str(c.fill.fgColor.rgb) if c.fill is not None and c.fill.fill_type == 'solid' else None
            celdas[c.coordinate] = (repr(c.value), c.number_format, relleno, bool(c.protection.locked),
                                    bool(c.alignment.wrap_text))
        fuera[ws.title] = {
            'celdas': celdas,
            'merges': sorted(str(m) for m in ws.merged_cells.ranges),
            'dv': sorted('%s:%s:%s:%s:%s:%s:%s' % (dv.type, dv.formula1, dv.formula2, dv.sqref, dv.error,
                                                   dv.errorTitle, dv.prompt)
                         for dv in ws.data_validations.dataValidation),
            'cf': sorted('%s:%s:%s' % (cf.sqref, [r.formula for r in cf.rules], [r.text for r in cf.rules])
                         for cf in ws.conditional_formatting),
            'prot': (bool(ws.protection.sheet), ws.protection.password), 'area': str(ws.print_area),
            'papel': ws.page_setup.paperSize, 'pie': ws.oddFooter.center.text,
            'alto': sorted((k, d.height) for k, d in ws.row_dimensions.items() if d.height),
            'ancho': sorted((k, d.width) for k, d in ws.column_dimensions.items() if d.width),
        }
    return fuera


def diferencias(a, b, fichero):
    out = []
    for k in ('_props', '_graficos'):
        if a[k] != b[k]:
            out.append('%s: %s' % (fichero, k))
    for hoja in sorted((set(a) | set(b)) - {'_props', '_graficos'}):
        if hoja not in a or hoja not in b:
            out.append('%s:%s: hoja solo en una pasada' % (fichero, hoja))
            continue
        for k in ('merges', 'dv', 'cf', 'prot', 'area', 'papel', 'pie', 'alto', 'ancho'):
            if a[hoja][k] != b[hoja][k]:
                out.append('%s:%s: cambia %s' % (fichero, hoja, k))
        ca, cb = a[hoja]['celdas'], b[hoja]['celdas']
        for coord in sorted(set(ca) | set(cb)):
            if ca.get(coord) != cb.get(coord):
                out.append('%s:%s!%s: %s → %s' % (fichero, hoja, coord, ca.get(coord), cb.get(coord)))
    return out


# ==========================================================================
# inject_cache (con IRR) y verificación de la caché
# ==========================================================================
def inject_cache(carpeta, log=print):
    fuera = []
    for fname_es in mapas.LIBROS_ES:
        termica()
        en = mapas.FICHEROS[fname_es]
        p = os.path.join(carpeta, en)
        try:
            nf, inj, fail = inyectar_cache(p)
            linea = '%s: formulas=%d cache_inyectado=%d fallos_pycel=%d' % (en, nf, inj, fail)
            ok = True
        except Exception as e:                                # noqa: BLE001
            linea, ok, fail = '%s: EXCEPCIÓN %r' % (en, e), False, -1
        log('    ' + linea)
        fuera.append({'fichero': en, 'salida': linea, 'ok': ok, 'fallos_pycel': fail})
    p07 = os.path.join(carpeta, FICHERO_07)
    irr = cachear_irr(p07)
    log('    cachear_irr %s: flujos %s · TIR Python %r · caché %r' % (FICHERO_07, irr['flujos'], irr['tir_python'],
                                                                     irr['cache_pycel']))
    return fuera, irr


def verificar_cache(carpeta):
    out = OrderedDict()
    for fname_es in mapas.LIBROS_ES:
        p = os.path.join(carpeta, mapas.FICHEROS[fname_es])
        wf, wv = openpyxl.load_workbook(p), openpyxl.load_workbook(p, data_only=True)
        n = con_valor = 0
        for ws in wf.worksheets:
            for c in ws._cells.values():
                if c.data_type == 'f':
                    n += 1
                    if wv[ws.title][c.coordinate].value is not None:
                        con_valor += 1
        out[mapas.FICHEROS[fname_es]] = (n, con_valor)
    return out


def carpeta_dryrun(arg):
    if arg:
        return os.path.abspath(arg)
    s = os.environ.get('CLAUDE_SCRATCHPAD')
    if s:
        return os.path.join(os.path.abspath(s), 'financial-dryrun')
    return os.path.join(REPO, '.work', mapas.SLUG, 'financial-dryrun')


def main():
    ap = argparse.ArgumentParser(description='Restaurant Financial Plan Kit Pro (EN): montaje de los 10 xlsx')
    ap.add_argument('--dry-run', action='store_true', help='escribe en el scratchpad (o --salida), no en dl/')
    ap.add_argument('--salida', default=None, help='carpeta del dry-run')
    ap.add_argument('--mes', default=None, help='mes de la línea de versión D20 (por defecto, el actual)')
    ap.add_argument('--json', default=None, help='informe JSON')
    ap.add_argument('--idempotencia', action='store_true', help='2.ª construcción en un clon y 0 diferencias')
    ap.add_argument('--parcial', action='store_true', help='(solo --dry-run) deja en ES lo que no tenga traducción')
    ap.add_argument('--sin-cache', action='store_true', help='(depuración) no corre inject_cache')
    args = ap.parse_args()
    if args.parcial and not args.dry_run:
        raise SystemExit('--parcial solo vale con --dry-run')

    mes = args.mes or mes_por_defecto()
    if mes not in MESES:
        raise SystemExit('--mes debe ser un mes en inglés: %s' % ', '.join(MESES))
    carpeta = carpeta_dryrun(args.salida) if args.dry_run else DESTINO_REAL
    if os.path.abspath(carpeta) == os.path.abspath(ORIGEN):
        raise SystemExit('ABORTADO: la salida no puede ser la carpeta ES')

    t0 = time.time()
    try:
        datos = cargar_datos(args.parcial)
        print('== 1/3 · montaje de los 10 libros → %s (mes %s)%s' % (
            carpeta, mes, ' · PARCIAL: %d cadenas sin traducción' % len(datos['faltan']) if args.parcial else ''),
            flush=True)
        informe = construir(carpeta, datos, mes)
    except Aborta as e:
        raise SystemExit('ABORTADO: %s' % e)
    sobran = sorted(set(os.listdir(carpeta)) - set(mapas.FICHEROS.values()))
    if sobran:
        print('  aviso: ficheros ajenos en la salida: %s' % sobran)

    idem = {'ejecutada': False}
    if args.idempotencia:
        print('== 2/3 · idempotencia: 2.ª construcción en un clon', flush=True)
        clon = tempfile.mkdtemp(prefix='financial-idem-')
        construir(clon, datos, mes, log=lambda *_: None)
        difs = []
        for n in mapas.FICHEROS.values():
            difs += diferencias(digest(os.path.join(carpeta, n)), digest(os.path.join(clon, n)), n)
        shutil.rmtree(clon)
        idem = {'ejecutada': True, 'diferencias': len(difs), 'detalle': difs[:30]}
        print('  diferencias 1.ª vs 2.ª construcción: %d' % len(difs), flush=True)
        for d in difs[:10]:
            print('    ' + d)

    print('== 3/3 · inject_cache (AL FINAL, sobre los 10, IRR incluida)', flush=True)
    cache, irr = ([], None) if args.sin_cache else inject_cache(carpeta)
    ver = verificar_cache(carpeta) if cache else {}
    for f, (n, v) in ver.items():
        print('    data_only %-50s %4d fórmulas · %4d con valor' % (f, n, v))

    fallos = []
    if idem.get('diferencias'):
        fallos.append('idempotencia: %d diferencias' % idem['diferencias'])
    fallos += ['inject_cache %s' % c['salida'] for c in cache if not c['ok'] or c['fallos_pycel'] != 0]
    salida = OrderedDict([('producto', mapas.SLUG), ('version', VERSION),
                          ('fecha', datetime.datetime.now().isoformat(timespec='seconds')),
                          ('modo', 'dry-run' if args.dry_run else 'real'), ('carpeta', carpeta), ('mes', mes),
                          ('libros', informe), ('idempotencia', idem), ('inject_cache', cache), ('irr', irr),
                          ('cache_data_only', ver), ('fallos', fallos),
                          ('segundos', round(time.time() - t0, 1))])
    if args.json:
        with open(args.json, 'w', encoding='utf-8') as fh:
            json.dump(salida, fh, ensure_ascii=False, indent=1, default=str)
        print('informe → %s' % args.json)
    if fallos:
        print('\nFALLOS:\n  ' + '\n  '.join(fallos))
        sys.exit(1)
    print('\nOK: 10 xlsx en %s (%.0f s)' % (carpeta, time.time() - t0))


if __name__ == '__main__':
    try:
        main()
    except Aborta as e:
        raise SystemExit('ABORTADO: %s' % e)
