#!/usr/bin/env python3
"""
extraer_textos.py — Restaurant Financial Plan Kit Pro (EN) · cimientos de la F2-EN.

Copia adaptada de `staff-kit/extraer_textos.py` (sesión Claude Code, 3-oct-2026), con la lectura de gráficos de
`recipe-costing-kit/extraer_textos.py`. Lee los 10 xlsx PUBLICADOS del Kit Plan Financiero v2.0
(`astro-site/public/dl/kit-plan-financiero/`, SOLO LECTURA) y escribe, junto a este script (o en --salida):

  · censo_es.json   por libro y hoja, lo que contarán los gates de la F2-EN: fórmulas (huella sha1 y
                    PATRONES normalizados con su recuento), referencias a hojas, literales en fórmulas,
                    en CF y en DV personalizadas, DV, CF (con sus tokens), GRÁFICOS (del ZIP: hoja, tipo,
                    título, referencias, textos de caché), merges, protección, paneles, áreas y títulos de
                    impresión, papel, cabeceras/pies, formatos (con €), fechas, números escritos a mano, textos
                    con €, con m² o con normativa/fiscalidad española, celdas verdes y desbloqueadas, docProps.
  · textos_es.json  cadenas de texto únicas con su contexto, TRATAMIENTO, PISTAS de la SPEC y GRUPO de
                    traducción (G1-G2 para los dos subagentes; GM = no se traduce).

    python3 extraer_textos.py [--origen DIR] [--salida DIR] [--md]

No traduce nada (regla 1bis: los textos EN los escriben subagentes Anthropic, nunca bridge.py).
Térmica: un solo proceso, en serie, sin builds ni navegador (10 libros, < 15 s).

Tratamientos (la cadena toma el más «fuerte»: traducir > localizar > mapa-hoja > mapa-clave > regenerar):
  traducir    texto libre → subagente (G1-G2), con pistas de la SPEC.
  localizar   dato de ejemplo → subagente, con SPEC §5 (este kit no tiene zonas: sus ejemplos son números).
  mapa-hoja   texto igual a una pestaña del kit (o a un mes) → `mapas.HOJAS`.
  mapa-clave  ítem de DV o literal de fórmula/CF del libro (o celda igual a uno) → `mapas.CLAVES`.
  regenerar   lo escribe `aplicar_en.py`: versión, línea de marca, pie, docProps, título interior,
              `mapas.FIJOS` (cadenas con decisión de mercado) y `mapas.POR_CELDA`.
"""
import argparse
import datetime
import hashlib
import json
import os
import posixpath
import re
import subprocess
import sys
import zipfile
from collections import OrderedDict, defaultdict

import openpyxl
from openpyxl.styles.numbers import is_date_format
from openpyxl.utils import get_column_letter, range_boundaries

AQUI = os.path.dirname(os.path.abspath(__file__))
REPO = os.path.abspath(os.path.join(AQUI, '..', '..', '..'))
ORIGEN = os.path.join(REPO, 'astro-site', 'public', 'dl', 'kit-plan-financiero')
sys.path.insert(0, AQUI)
import mapas                                                     # noqa: E402

VERDE = 'E8F5E9'
INVARIABLES = {'OK', 'N/A', 'AI Chef Pro', '© 2026 AI Chef Pro · aichef.pro', 'EBITDA', 'Marketing'}
RX_VERSION = re.compile(r'^Versi[óo]n \d+\.\d+ · ')
RX_LITERAL = re.compile(r'"((?:[^"]|"")*)"')
RX_REF_HOJA = re.compile(r"(?:'((?:[^']|'')+)'|([^\W\d][\w.]*))!(?=\$?[A-Za-z]{1,3}\$?\d|\$?[A-Za-z]{1,3}:)")
RX_CITA = re.compile(r"«([^»]{1,40})»|'([^']{1,40})'|\"([^\"]{1,40})\"")
RX_NORMA = re.compile(r'\bIVA\b|IRPF|Seguridad Social|\bSS\b|RETA|\bCCC\b|[Mm]od(?:elo|\.) ?\d{3}|AEAT|Hacienda|\bICO\b|'
                      r'ENISA|\bLIS\b|[Ii]mpuesto (?:de|sobre) Sociedades|\bSL\b|[Nn]otar|Registro Mercantil|'
                      r'\bNIF\b|\bCIF\b|OEPM|RGSEAA|SGAE|Veri\*factu|[Cc]onvenio|[Gg]estoría|\bTIN\b|\bTAE\b|'
                      r'\bSGR\b|[Aa]utónom|autonómic|\bart\.|[Cc]arencia|[Aa]val\b|[Ff]ianza|[Tt]raspaso')
RX_MONEDA = re.compile(r'€|\bEUR\b|euros?\b')
RX_M2 = re.compile(r'm²')
RX_COMPARA = re.compile(r'[<>]=?\s*-?\d')
RX_HORA_FMT = re.compile(r'h{1,2}:mm', re.I)

# Zonas de DATOS DE EJEMPLO por hoja (fila1, fila2, columnas): ninguna — los ejemplos del kit son números (D14)
ZONAS_EJEMPLO = {}
ZONAS_REGENERAR = {}

PISTAS = [
    (RX_NORMA.pattern, 'D8-D13: la fiscalidad/administración española se sustituye por la referencia US de SPEC §3 '
                       '(sales tax, payroll taxes IRS, income tax, SBA 7(a), state/city permits) con nota UK (VAT, '
                       'corporation tax, PAYE) solo donde la SPEC la pone; nunca se traduce un modelo, una ley o un '
                       'organismo español. Lo que no tiene equivalente se reescribe o se quita (SPEC §4)'),
    (RX_MONEDA.pattern, 'D18: sin símbolo de moneda («in your currency»); quitar «(€)», «€/mes» → «per month»'),
    (r'm²', 'D14: m² → sq ft (superficie de sala y ventas por sq ft)'),
    (r'[Ss]in IVA|[Cc]on IVA|IVA incl', 'D8: «sin IVA» → «excl. sales tax»; «con IVA / IVA incl.» → «incl. sales tax»; '
                                        'en el 03 la caja va CON el impuesto cobrado'),
    (r'[Cc]ubiertos?|[Tt]icket medio|[Cc]omensal', 'Glosario: cubiertos → «covers»; ticket medio → «average check»'),
    (r'[Ff]acturaci|[Ii]ngresos|[Vv]entas', 'Glosario: facturación/ingresos → «revenue» o «sales»'),
    (r'[Pp]ersonal|[Nn]ómina', 'Glosario: personal → «labor»; nóminas → «payroll»'),
    (r'[Ss]ala\b|[Bb]arra|[Cc]omedor', 'Glosario: comedor → «dine-in»; barra → «bar»; sala → «dining room»'),
    (r'TPV|[Dd]atáfono', 'Glosario: TPV → POS; datáfono → card terminal'),
    (r'[Aa]mortizaci', 'D10/D12: amortización del inmovilizado → «depreciation»; amortización de un préstamo → '
                       '«principal repayment»; cuadro de amortización → «amortization schedule»'),
    (r'[Pp]réstamo|[Bb]anco|[Ee]ntidad', 'D12-D13: banco → «lender»; el informe va a «lenders and investors»'),
    (r'TIR|VAN|[Pp]ayback', 'Glosario: TIR → IRR; VAN → NPV; payback → payback period'),
    (r'A4', 'D19: papel US Letter'),
    (r'Revisar →|[Dd]esproteger|Archivo →', 'Rutas de Excel en inglés: «Review → Unprotect Sheet»'),
    (r'\b0\d\b|BONUS-0\d|\b01b\b', 'D5: cita de plantilla → «file NN» + nombre corto (mapas.NOMBRES_CITA), un solo nivel'),
    (r"[«»]|'[A-ZÁÉÍÓÚ]", 'D6: pestañas y valores citados entre comillas dobles rectas ("Monthly Cash Flow")'),
    (r'EJEMPLO|[Ee]jemplo|orientativo', '«valores de ejemplo» → «sample values»; «orientativo» → «example»'),
    (r'\d,\d', 'Formato US: punto decimal y coma de miles'),
    (r'p\.p\.', 'D18: puntos porcentuales → «pts» (percentage points)'),
    (r'antes este informe|hasta hoy|[Vv]enía de', 'D20: la EN nace en 2.0 — se explica el porqué sin contar la historia de '
                                                'versiones («antes este informe pedía…», «hasta hoy…», «venía de…»)'),
    (r'[Cc]arencia', 'D12: carencia → «interest-only period»'),
    (r'[Ii]mpuesto (?:de|sobre) Sociedades|\bIS\b|BAI', 'D12: Impuesto sobre Sociedades → «income tax»; BAI → «EBT '
                                                       '(earnings before tax)»'),
    (r'Kit Plan Financiero', 'D2: «Restaurant Financial Plan Kit Pro»'),
    (r'aichef\.pro/kit-plan-financiero', 'D20: ' + mapas.URL_PRODUCTO),
]


# --------------------------------------------------------------------------
# utilidades
# --------------------------------------------------------------------------
def sha256(path):
    h = hashlib.sha256()
    with open(path, 'rb') as fh:
        for b in iter(lambda: fh.read(65536), b''):
            h.update(b)
    return h.hexdigest()


def corto(fichero):
    return mapas.corto_de(fichero)


def tiene_letra(s):
    return any(ch.isalpha() for ch in s)


def literales(formula):
    return [m.group(1).replace('""', '"') for m in RX_LITERAL.finditer(formula or '')]


def refs_hoja(formula):
    sin_lit = RX_LITERAL.sub('""', formula or '')
    return [(m.group(1) or m.group(2)).replace("''", "'") for m in RX_REF_HOJA.finditer(sin_lit)]


def celdas_de(sqref):
    s = set()
    for rng in str(sqref).split():
        c1, r1, c2, r2 = range_boundaries(rng)
        for r in range(r1, r2 + 1):
            for c in range(c1, c2 + 1):
                s.add((r, c))
    return s


def en_zona(zonas, clave, cel):
    z = zonas.get(clave)
    if not z:
        return None
    if z[0] <= cel.row <= z[1] and get_column_letter(cel.column) in z[2]:
        return z
    return None


def items_dv(f1):
    return f1.strip('"').split(',') if f1 and f1.startswith('"') else []


def claves_de_hoja(ws):
    """Ítems de DV de lista y literales de fórmula/CF/DV de la hoja: lo que una celda puede «ser»."""
    s = set()
    for dv in ws.data_validations.dataValidation:
        s.update(items_dv(dv.formula1))
        if dv.type == 'custom':
            s.update(literales(dv.formula1))
    for row in ws.iter_rows():
        for c in row:
            if isinstance(c.value, str) and c.data_type == 'f':
                s.update(literales(c.value))
    for rng, reglas in ws.conditional_formatting._cf_rules.items():
        for rg in reglas:
            for f in rg.formula or []:
                s.update(literales(f))
    return s


def clasificar_celda(fichero, ws, cel, texto, hojas_libro, claves_hoja):
    f = corto(fichero)
    if (f, ws.title, cel.coordinate) in mapas.POR_CELDA:
        return 'regenerar', 'rótulo fijado (mapas.POR_CELDA)'
    if ws.title == 'Instrucciones' and cel.coordinate == 'B2':
        return 'regenerar', 'D5 título interior (mapas.titulo_interior)'
    if RX_VERSION.match(texto):
        return 'regenerar', 'D20 línea de versión'
    if texto == mapas.MARCA_ES:
        return 'regenerar', 'D2 línea de marca «%s»' % mapas.MARCA_EN
    if texto in mapas.FIJOS:
        return 'regenerar', 'decisión de mercado (mapas.FIJOS)'
    if texto in mapas.HOJAS:
        return 'mapa-hoja', 'celda igual a una pestaña o a un mes'
    if texto in claves_hoja and texto in mapas.CLAVES:
        return 'mapa-clave', 'valor de DV/literal del libro'
    if en_zona(ZONAS_EJEMPLO, (f, ws.title), cel):
        return 'localizar', 'dato de ejemplo (SPEC §5)'
    return 'traducir', None


def _xml_attr(tag_xml, attr):
    m = re.search(r'\b' + attr + r'="([^"]*)"', tag_xml)
    return m.group(1) if m else None


def _resolver(base_parte, target):
    if target.startswith('/'):
        return target.lstrip('/')
    return posixpath.normpath(posixpath.join(posixpath.dirname(base_parte), target))


def graficos(path):
    """Gráficos del libro leídos del ZIP: hoja ancla, tipo, título, refs y textos de caché (copia del recipe-kit)."""
    out = []
    with zipfile.ZipFile(path) as z:
        nombres = set(z.namelist())
        if not any(n.startswith('xl/charts/') for n in nombres):
            return out
        wbx = z.read('xl/workbook.xml').decode('utf-8')
        rels = z.read('xl/_rels/workbook.xml.rels').decode('utf-8')
        rid2t = {_xml_attr(t, 'Id'): _xml_attr(t, 'Target') for t in re.findall(r'<Relationship\b[^>]*>', rels)}
        hoja_de_parte = {}
        for t in re.findall(r'<sheet\b[^>]*>', wbx):
            parte = _resolver('xl/workbook.xml', rid2t[_xml_attr(t, 'r:id')])
            hoja_de_parte[parte] = _xml_attr(t, 'name').replace('&amp;', '&')
        ancla = {}
        for parte, hoja in hoja_de_parte.items():
            srel = posixpath.join(posixpath.dirname(parte), '_rels', posixpath.basename(parte) + '.rels')
            if srel not in nombres:
                continue
            for t in re.findall(r'<Relationship\b[^>]*>', z.read(srel).decode('utf-8')):
                if _xml_attr(t, 'Type').endswith('/drawing'):
                    dparte = _resolver(parte, _xml_attr(t, 'Target'))
                    drel = posixpath.join(posixpath.dirname(dparte), '_rels', posixpath.basename(dparte) + '.rels')
                    if drel in nombres:
                        for t2 in re.findall(r'<Relationship\b[^>]*>', z.read(drel).decode('utf-8')):
                            if _xml_attr(t2, 'Type').endswith('/chart'):
                                ancla[_resolver(dparte, _xml_attr(t2, 'Target'))] = hoja
        for parte in sorted(n for n in nombres if re.fullmatch(r'xl/charts/chart\d+\.xml', n)):
            x = z.read(parte).decode('utf-8')
            titulo = None
            mt = re.search(r'<(?:c:)?title>(.*?)</(?:c:)?title>', x, re.S)
            if mt:
                titulo = ''.join(re.findall(r'<a:t>([^<]*)</a:t>', mt.group(1))) or None
            ejes = [''.join(re.findall(r'<a:t>([^<]*)</a:t>', m)) for m in
                    re.findall(r'<(?:c:)?(?:valAx|catAx)>.*?<(?:c:)?title>(.*?)</(?:c:)?title>', x, re.S)]
            refs = [r.replace('&apos;', "'").replace('&amp;', '&') for r in re.findall(r'<(?:c:)?f>([^<]*)</(?:c:)?f>', x)]
            out.append(OrderedDict([
                ('parte', parte), ('hoja', ancla.get(parte)), ('tipos', re.findall(r'<(?:c:)?(\w+Chart)>', x)),
                ('titulo', titulo), ('titulos_eje', [e for e in ejes if e]),
                ('series', len(re.findall(r'<(?:c:)?ser>', x))),
                ('refs', refs), ('hojas_citadas', sorted({h for r in refs for h in refs_hoja(r)})),
                ('textos_cache', re.findall(r'<(?:c:)?v>([^<]*[A-Za-zÁ-ú][^<]*)</(?:c:)?v>', x)),
            ]))
    return out


# --------------------------------------------------------------------------
# censo de un libro
# --------------------------------------------------------------------------
def censar_libro(path, fichero):
    wb = openpyxl.load_workbook(path)
    hojas = list(wb.sheetnames)
    ocurr = []
    info = OrderedDict()
    info['corto'] = corto(fichero)
    info['sha256'] = sha256(path)
    info['nombre_en'] = mapas.FICHEROS[fichero]
    info['titulo_en'] = mapas.TITULOS[mapas.FICHEROS[fichero]]
    info['hojas'] = hojas
    info['hojas_en'] = [mapas.HOJAS.get(h) for h in hojas]
    info['nombres_definidos'] = sorted(getattr(wb.defined_names, 'keys', lambda: [])())
    pr = wb.properties
    info['docprops'] = OrderedDict((k, getattr(pr, k)) for k in
                                   ('creator', 'title', 'subject', 'keywords', 'description', 'category',
                                    'lastModifiedBy'))
    info['hojas_detalle'] = OrderedDict()
    for h in hojas:
        ocurr.append((h, {'f': corto(fichero), 'h': h, 't': 'hoja', 'tr': 'mapa-hoja', 'det': 'nombre de pestaña'}))

    fx_hoja_total = 0
    claves_libro = set()
    for ws in wb.worksheets:
        claves_libro |= claves_de_hoja(ws)
    for ws in wb.worksheets:
        hd = OrderedDict()
        hd['nombre_en'] = mapas.HOJAS.get(ws.title)
        hd['dimension'] = ws.dimensions
        hd['max_fila'] = ws.max_row
        hd['max_col'] = ws.max_column
        claves_hoja = claves_libro
        n_txt = n_num = n_fx = fx_con_hoja = 0
        refs = defaultdict(int)
        lits = defaultdict(list)
        formatos = defaultdict(int)
        fechas, horas, numeros, m2, norma, moneda, fmt_euro = [], [], [], [], [], [], []
        verdes = verdes_desbl = desbl = 0
        patrones = OrderedDict()
        huella = hashlib.sha1()
        for row in ws.iter_rows():
            for cel in row:
                if cel.__class__.__name__ == 'MergedCell':
                    continue
                v = cel.value
                fmt = cel.number_format
                fid = getattr(cel._style, 'numFmtId', None)
                if is_date_format(fmt):
                    d = OrderedDict([('celda', cel.coordinate), ('numFmtId', fid), ('formato', fmt),
                                     ('vacia', v is None)])
                    (horas if RX_HORA_FMT.search(fmt) and not re.search(r'[dy]', fmt, re.I) else fechas).append(d)
                if fmt and fmt != 'General':
                    formatos[fmt] += 1
                if fmt and '€' in fmt:
                    fmt_euro.append(cel.coordinate)
                fill = cel.fill.fgColor.rgb if cel.fill is not None and cel.fill.fgColor is not None else None
                bloq = cel.protection is None or cel.protection.locked is not False
                if not bloq:
                    desbl += 1
                if isinstance(fill, str) and fill.upper().endswith(VERDE):
                    verdes += 1
                    if not bloq:
                        verdes_desbl += 1
                if v is None:
                    continue
                if isinstance(v, str) and cel.data_type == 'f':
                    n_fx += 1
                    huella.update('{}={}\n'.format(cel.coordinate, v).encode('utf-8'))
                    for lit in literales(v):
                        lits[lit].append(cel.coordinate)
                    rh = refs_hoja(v)
                    if rh:
                        fx_con_hoja += 1
                    for nom in rh:
                        refs[nom] += 1
                    p = mapas.patron(v)
                    if p not in patrones:
                        patrones[p] = OrderedDict([('patron', p), ('n', 0), ('celdas', []),
                                                   ('cifras', bool(RX_COMPARA.search(RX_LITERAL.sub('""', v))))])
                    patrones[p]['n'] += 1
                    patrones[p]['celdas'].append(cel.coordinate)
                    continue
                if isinstance(v, str):
                    n_txt += 1
                    if RX_M2.search(v):
                        m2.append(cel.coordinate)
                    if RX_NORMA.search(v):
                        norma.append(cel.coordinate)
                    if RX_MONEDA.search(v):
                        moneda.append(cel.coordinate)
                    tr, det = clasificar_celda(fichero, ws, cel, v, hojas, claves_hoja)
                    d = {'f': corto(fichero), 'h': ws.title, 'c': cel.coordinate, 't': 'celda', 'tr': tr}
                    if det:
                        d['det'] = det
                    ocurr.append((v, d))
                else:
                    n_num += 1
                    if ws.title != 'Instrucciones':
                        val = v.isoformat() if hasattr(v, 'isoformat') else v
                        numeros.append([cel.coordinate, val, fmt])
        fx_hoja_total += fx_con_hoja
        for p in patrones.values():
            cs = p['celdas']
            p['celdas'] = cs[0] + ('..' + cs[-1] if len(cs) > 1 else '')
        hd['celdas_texto'] = n_txt
        hd['celdas_numero'] = n_num
        hd['celdas_formula'] = n_fx
        hd['formulas_con_hoja'] = fx_con_hoja
        hd['refs_hoja'] = OrderedDict(sorted(refs.items()))
        hd['literales'] = OrderedDict((k, lits[k]) for k in sorted(lits))
        hd['literales_recuento'] = OrderedDict((k, len(lits[k])) for k in sorted(lits))
        hd['huella_formulas_sha1'] = huella.hexdigest() if n_fx else None
        hd['patrones_formula'] = list(patrones.values())
        hd['fechas'] = fechas
        hd['horas'] = horas
        hd['numeros'] = numeros
        hd['formatos'] = OrderedDict(sorted(formatos.items(), key=lambda kv: -kv[1]))
        hd['formatos_euro'] = fmt_euro
        hd['textos_m2'] = m2
        hd['textos_normativa_es'] = norma
        hd['textos_moneda'] = moneda
        hd['verdes'] = verdes
        hd['verdes_desbloqueadas'] = verdes_desbl
        hd['desbloqueadas'] = desbl
        hd['columnas_ocultas'] = sorted(k for k, d in ws.column_dimensions.items() if d.hidden)
        hd['filas_ocultas'] = sorted(k for k, d in ws.row_dimensions.items() if d.hidden)
        hd['merges'] = sorted(str(m) for m in ws.merged_cells.ranges)
        hd['paneles'] = ws.freeze_panes
        p = ws.protection
        hd['proteccion'] = OrderedDict([('hoja', bool(p.sheet)), ('password', bool(p.password)),
                                        ('formatColumns', p.formatColumns), ('formatRows', p.formatRows),
                                        ('insertRows', p.insertRows), ('deleteRows', p.deleteRows),
                                        ('sort', p.sort), ('autoFilter', p.autoFilter)])
        hd['autofiltro'] = ws.auto_filter.ref
        hd['area_impresion'] = ws.print_area
        hd['titulos_impresion'] = OrderedDict([('filas', ws.print_title_rows), ('columnas', ws.print_title_cols)])
        ps = ws.page_setup
        hd['papel'] = OrderedDict([('paperSize', ps.paperSize), ('orientation', ps.orientation),
                                   ('fitToWidth', ps.fitToWidth), ('fitToHeight', ps.fitToHeight),
                                   ('fitToPage', bool(ws.sheet_properties.pageSetUpPr and
                                                      ws.sheet_properties.pageSetUpPr.fitToPage))])
        hf = OrderedDict()
        for nombre in ('oddHeader', 'oddFooter', 'evenHeader', 'evenFooter', 'firstHeader', 'firstFooter'):
            obj = getattr(ws, nombre)
            for parte in ('left', 'center', 'right'):
                txt = getattr(obj, parte).text
                if txt:
                    hf['{}.{}'.format(nombre, parte)] = txt
                    ocurr.append((txt, {'f': corto(fichero), 'h': ws.title, 't': 'cabecera-pie',
                                        'c': '{}.{}'.format(nombre, parte), 'tr': 'regenerar',
                                        'det': 'D18 pie «%s»' % mapas.FOOTER_EN}))
        hd['cabeceras_pies'] = hf
        dvs = []
        for dv in ws.data_validations.dataValidation:
            dvs.append(OrderedDict([
                ('sqref', str(dv.sqref)), ('type', dv.type), ('operator', dv.operator),
                ('formula1', dv.formula1), ('formula2', dv.formula2), ('allow_blank', dv.allow_blank),
                ('showErrorMessage', dv.showErrorMessage), ('showInputMessage', dv.showInputMessage),
                ('errorStyle', dv.errorStyle),
                ('errorTitle', dv.errorTitle), ('error', dv.error),
                ('promptTitle', dv.promptTitle), ('prompt', dv.prompt),
                ('n_celdas', len(celdas_de(dv.sqref))),
                ('cita_hojas', sorted(set(refs_hoja(dv.formula1 or '')))),
                ('literales', literales(dv.formula1) if dv.type == 'custom' else []),
            ]))
            base = {'f': corto(fichero), 'h': ws.title, 'c': str(dv.sqref)}
            for campo, t in (('errorTitle', 'dv-error-titulo'), ('error', 'dv-error'),
                             ('promptTitle', 'dv-prompt-titulo'), ('prompt', 'dv-prompt')):
                txt = getattr(dv, campo)
                if txt:
                    ocurr.append((txt, dict(base, t=t, tr='traducir')))
            for it in items_dv(dv.formula1):
                if it == '':
                    continue
                d = dict(base, t='dv-item', tr='mapa-clave', det='mapas.CLAVES')
                ocurr.append((it, d))
        hd['dv'] = dvs
        cfs = []
        lits_cf = defaultdict(int)
        for rng, reglas in ws.conditional_formatting._cf_rules.items():
            for rg in reglas:
                fs = list(rg.formula or [])
                cfs.append(OrderedDict([
                    ('sqref', str(rng.sqref)), ('type', rg.type), ('operator', rg.operator),
                    ('text', getattr(rg, 'text', None)), ('formula', fs), ('priority', rg.priority),
                    ('stopIfTrue', rg.stopIfTrue), ('cita_hojas', sorted({h for f in fs for h in refs_hoja(f)})),
                    ('tokens', [t for f in fs for t in literales(f)]),
                    ('fill', getattr(getattr(getattr(rg.dxf, 'fill', None), 'fgColor', None), 'rgb', None)
                     if rg.dxf is not None else None),
                ]))
                for f in fs:
                    for lit in literales(f):
                        lits_cf[lit] += 1
        hd['cf'] = cfs
        hd['literales_cf'] = OrderedDict(sorted(lits_cf.items()))
        info['hojas_detalle'][ws.title] = hd

    info['graficos'] = graficos(path)
    with zipfile.ZipFile(path) as z:
        info['imagenes'] = sorted(n for n in z.namelist() if n.startswith('xl/media/'))
    for g in info['graficos']:
        for campo, txts in (('grafico-titulo', [g['titulo']] if g['titulo'] else []), ('grafico-eje', g['titulos_eje'])):
            for t in txts:
                d = {'f': corto(fichero), 'h': g['hoja'], 't': campo, 'c': g['parte'], 'tr': 'traducir'}
                if t in mapas.FIJOS:
                    d.update(tr='regenerar', det='decisión de mercado (mapas.FIJOS)')
                ocurr.append((t, d))
    for campo, v in info['docprops'].items():
        if not isinstance(v, str) or not v or campo == 'creator':
            continue
        ocurr.append((v, {'f': corto(fichero), 't': 'docprops', 'c': campo, 'tr': 'regenerar',
                          'det': 'D18 docProps'}))

    hdv = list(info['hojas_detalle'].values())
    lit_tot, litcf_tot = defaultdict(int), defaultdict(int)
    for x in hdv:
        for k, v in x['literales_recuento'].items():
            lit_tot[k] += v
        for k, v in x['literales_cf'].items():
            litcf_tot[k] += v
    info['totales'] = OrderedDict([
        ('hojas', len(hojas)),
        ('celdas_texto', sum(x['celdas_texto'] for x in hdv)),
        ('celdas_formula', sum(x['celdas_formula'] for x in hdv)),
        ('patrones_formula', sum(len(x['patrones_formula']) for x in hdv)),
        ('formulas_con_hoja', fx_hoja_total),
        ('dv', sum(len(x['dv']) for x in hdv)),
        ('dv_lista_literal', sum(1 for x in hdv for d in x['dv'] if (d['formula1'] or '').startswith('"'))),
        ('dv_cita_hoja', sum(1 for x in hdv for d in x['dv'] if d['cita_hojas'])),
        ('dv_celdas', sum(d['n_celdas'] for x in hdv for d in x['dv'])),
        ('cf', sum(len(x['cf']) for x in hdv)),
        ('graficos', len(info['graficos'])),
        ('series_grafico', sum(g['series'] for g in info['graficos'])),
        ('imagenes', len(info['imagenes'])),
        ('merges', sum(len(x['merges']) for x in hdv)),
        ('hojas_protegidas', sum(1 for x in hdv if x['proteccion']['hoja'])),
        ('hojas_con_password', sum(1 for x in hdv if x['proteccion']['password'])),
        ('areas_impresion', sum(1 for x in hdv if x['area_impresion'])),
        ('titulos_impresion', sum(1 for x in hdv if x['titulos_impresion']['filas'])),
        ('paneles', sum(1 for x in hdv if x['paneles'])),
        ('fechas', sum(len(x['fechas']) for x in hdv)),
        ('horas', sum(len(x['horas']) for x in hdv)),
        ('numeros_ejemplo', sum(len(x['numeros']) for x in hdv)),
        ('formatos_euro', sum(len(x['formatos_euro']) for x in hdv)),
        ('textos_m2', sum(len(x['textos_m2']) for x in hdv)),
        ('textos_normativa_es', sum(len(x['textos_normativa_es']) for x in hdv)),
        ('textos_moneda', sum(len(x['textos_moneda']) for x in hdv)),
        ('verdes', sum(x['verdes'] for x in hdv)),
        ('verdes_desbloqueadas', sum(x['verdes_desbloqueadas'] for x in hdv)),
        ('desbloqueadas', sum(x['desbloqueadas'] for x in hdv)),
        ('columnas_ocultas', sum(len(x['columnas_ocultas']) for x in hdv)),
        ('papel', sorted({x['papel']['paperSize'] for x in hdv if x['papel']['paperSize'] is not None})),
        ('literales', OrderedDict(sorted(lit_tot.items()))),
        ('literales_cf', OrderedDict(sorted(litcf_tot.items()))),
    ])
    return info, ocurr


# --------------------------------------------------------------------------
# cadenas, grupos y pistas
# --------------------------------------------------------------------------
PRIORIDAD = {'traducir': 0, 'localizar': 1, 'mapa-hoja': 2, 'mapa-clave': 3, 'regenerar': 4}
SUMABLES = ('hojas', 'celdas_texto', 'celdas_formula', 'patrones_formula', 'formulas_con_hoja', 'dv',
            'dv_lista_literal', 'dv_cita_hoja', 'dv_celdas', 'cf', 'graficos', 'series_grafico', 'imagenes',
            'merges', 'hojas_protegidas', 'hojas_con_password', 'areas_impresion', 'titulos_impresion', 'paneles',
            'fechas', 'horas', 'numeros_ejemplo', 'formatos_euro', 'textos_m2', 'textos_normativa_es',
            'textos_moneda', 'verdes', 'verdes_desbloqueadas', 'desbloqueadas', 'columnas_ocultas')


def pistas(texto):
    return [nota for rx, nota in PISTAS if re.search(rx, texto)]


def citas_hoja(texto, todas_hojas):
    out = []
    for m in RX_CITA.finditer(texto):
        span = next(g for g in m.groups() if g is not None)
        if span in todas_hojas and span not in [c['es'] for c in out]:
            out.append(OrderedDict([('es', span), ('en', mapas.HOJAS.get(span))]))
    return out


def palabras(s):
    return len(re.findall(r'\w+', s))


def dos_bloques(orden, peso):
    """Parte la lista de libros en 2 bloques contiguos minimizando el bloque más pesado (G1 lleva las comunes)."""
    mejor = None
    for i in range(1, len(orden)):
        bl = [orden[:i], orden[i:]]
        m = max(sum(peso[f] for f in b) + (peso['G1-comun'] if k == 0 else 0) for k, b in enumerate(bl))
        if mejor is None or m < mejor[0]:
            mejor = (m, bl)
    return mejor[1]


def main():
    ap = argparse.ArgumentParser(description=__doc__.split('\n')[1])
    ap.add_argument('--origen', default=ORIGEN)
    ap.add_argument('--salida', default=AQUI)
    ap.add_argument('--md', action='store_true', help='imprime la tabla §0 de F1-inventario-es.md')
    args = ap.parse_args()

    libros = list(mapas.LIBROS_ES)
    faltan = [f for f in libros if not os.path.exists(os.path.join(args.origen, f))]
    if faltan:
        sys.exit('faltan ficheros en {}: {}'.format(args.origen, faltan))
    sobran = sorted(f for f in os.listdir(args.origen) if f.endswith('.xlsx') and f not in libros)
    if sobran:
        sys.exit('xlsx no previstos en {}: {}'.format(args.origen, sobran))
    try:
        commit = subprocess.check_output(['git', '-C', REPO, 'rev-parse', '--short', 'HEAD'],
                                         stderr=subprocess.DEVNULL).decode().strip()
    except Exception:                                             # noqa: BLE001
        commit = None

    censo = OrderedDict()
    censo['meta'] = OrderedDict([
        ('generado', datetime.datetime.utcnow().replace(microsecond=0).isoformat() + 'Z'),
        ('script', 'scripts/productos-digitales/financial-kit/extraer_textos.py'),
        ('origen', os.path.relpath(args.origen, REPO)),
        ('commit', commit),
        ('nota', 'Censo del Kit Plan Financiero v2.0 PUBLICADO. Los gates de la F2-EN comparan contra '
                 'este fichero, nunca contra recuentos fijos. patrones_formula = fórmulas con las filas '
                 'normalizadas a {r} (mapas.patron); «cifras» = el patrón compara con un número (lo decide '
                 'mapas.FORMULAS_EN o mapas.PATRONES_SIN_CAMBIO).'),
    ])
    censo['libros'] = OrderedDict()
    todas = []
    for f in libros:
        info, oc = censar_libro(os.path.join(args.origen, f), f)
        censo['libros'][f] = info
        todas.extend(oc)
    T = OrderedDict((k, sum(censo['libros'][f]['totales'][k] for f in libros)) for k in SUMABLES)
    lit, litcf = defaultdict(int), defaultdict(int)
    for f in libros:
        for k, v in censo['libros'][f]['totales']['literales'].items():
            lit[k] += v
        for k, v in censo['libros'][f]['totales']['literales_cf'].items():
            litcf[k] += v
    T['literales'] = OrderedDict(sorted(lit.items(), key=lambda kv: -kv[1]))
    T['literales_cf'] = OrderedDict(sorted(litcf.items(), key=lambda kv: -kv[1]))
    T['papel'] = sorted({p for f in libros for p in censo['libros'][f]['totales']['papel']})
    censo['totales'] = T

    # ---- cadenas únicas
    todas_hojas = set(h for f in libros for h in censo['libros'][f]['hojas'])
    excl = OrderedDict()
    cad = OrderedDict()
    for texto, oc in todas:
        fijado = texto in mapas.FIJOS or (oc.get('f'), oc.get('h'), oc.get('c')) in mapas.POR_CELDA
        if (not tiene_letra(texto) and not RX_MONEDA.search(texto) and not fijado) or texto in INVARIABLES:
            excl[texto] = excl.get(texto, 0) + 1
            continue
        cad.setdefault(texto, []).append(oc)
    orden_f = [corto(f) for f in libros]
    cadenas = []
    for i, (texto, ocs) in enumerate(cad.items(), 1):
        fich = []
        for o in ocs:
            if o['f'] not in fich:
                fich.append(o['f'])
        trs = sorted({o['tr'] for o in ocs}, key=lambda t: PRIORIDAD[t])
        e = OrderedDict()
        e['id'] = 'c{:04d}'.format(i)
        e['es'] = texto
        e['grupo'] = None
        e['tratamiento'] = trs[0]
        if len(trs) > 1:
            e['tratamientos'] = trs
        if texto in mapas.HOJAS and 'mapa-hoja' in trs:
            e['en_mapa'] = mapas.HOJAS[texto]
        elif texto in mapas.CLAVES and 'mapa-clave' in trs:
            e['en_mapa'] = mapas.CLAVES[texto]
        e['n_ficheros'] = len(fich)
        e['n_apariciones'] = len(ocs)
        e['palabras'] = palabras(texto)
        e['con_cifra'] = bool(re.search(r'\d', texto))
        ch = citas_hoja(texto, todas_hojas)
        if ch:
            e['cita_hojas'] = ch
        if trs[0] in ('traducir', 'localizar'):
            pi = pistas(texto)
            if pi:
                e['pistas'] = pi
        e['sha1'] = hashlib.sha1(texto.encode('utf-8')).hexdigest()[:12]
        e['donde'] = ocs
        e['_ficheros'] = fich
        cadenas.append(e)

    # ---- grupos: G1-G2 (subagentes) por bloques contiguos de libros; comunes → G1; GM no se traduce
    peso = defaultdict(int)
    for e in cadenas:
        if e['tratamiento'] in ('traducir', 'localizar'):
            comun = e['n_ficheros'] >= 3
            e['_comun'] = comun
            peso['G1-comun' if comun else e['_ficheros'][0]] += e['palabras']
    bloques = dos_bloques(orden_f, peso)
    gid_de = {}
    for k, bl in enumerate(bloques, 1):
        for f in bl:
            gid_de[f] = 'G%d' % k
    for e in cadenas:
        if e['tratamiento'] in ('traducir', 'localizar'):
            e['grupo'] = 'G1' if e['_comun'] else gid_de[e['_ficheros'][0]]
            if e['_comun']:
                e['comun'] = True
        else:
            e['grupo'] = 'GM'
    grupos = []
    for k, bl in enumerate(bloques, 1):
        gid = 'G%d' % k
        es = [e for e in cadenas if e['grupo'] == gid]
        grupos.append(OrderedDict([
            ('id', gid),
            ('descripcion', ('Cadenas comunes a ≥ 3 libros (se traducen PRIMERO: fijan el glosario) + ' if k == 1 else '')
             + 'libros ' + ', '.join(bl)),
            ('libros', bl),
            ('n_cadenas', len(es)), ('n_localizar', sum(1 for e in es if e['tratamiento'] == 'localizar')),
            ('n_palabras', sum(e['palabras'] for e in es)),
        ]))
    gm = [e for e in cadenas if e['grupo'] == 'GM']
    por_tr = defaultdict(int)
    for e in gm:
        por_tr[e['tratamiento']] += 1
    grupos.append(OrderedDict([
        ('id', 'GM'),
        ('descripcion', 'NO se traduce: pestañas, meses y claves (mapas.py), FIJOS, POR_CELDA, versión, marca, pie, '
                        'docProps y título interior. ' +
                        ', '.join('{} {}'.format(v, k) for k, v in sorted(por_tr.items()))),
        ('n_cadenas', len(gm)), ('n_palabras', sum(e['palabras'] for e in gm)),
    ]))
    orden_g = {g['id']: k for k, g in enumerate(grupos)}
    cadenas.sort(key=lambda e: (orden_g[e['grupo']], int(e['id'][1:])))
    for e in cadenas:
        e.pop('_ficheros')
        e.pop('_comun', None)
    assert all(e['grupo'] for e in cadenas)

    textos = OrderedDict()
    textos['meta'] = OrderedDict([
        ('generado', censo['meta']['generado']),
        ('script', censo['meta']['script']),
        ('origen', censo['meta']['origen']),
        ('commit', commit),
        ('libros', OrderedDict((corto(f), OrderedDict([('es', f), ('en', mapas.FICHEROS[f]),
                                                      ('sha256', censo['libros'][f]['sha256'])]))
                               for f in libros)),
        ('n_cadenas', len(cadenas)),
        ('n_traducir', sum(1 for e in cadenas if e['tratamiento'] == 'traducir')),
        ('n_localizar', sum(1 for e in cadenas if e['tratamiento'] == 'localizar')),
        ('n_palabras_subagentes', sum(e['palabras'] for e in cadenas if e['grupo'] != 'GM')),
        ('n_apariciones', sum(e['n_apariciones'] for e in cadenas)),
        ('grupos', grupos),
        ('excluidas', OrderedDict([('cadenas', excl),
                                   ('regla', 'Fuera: cadenas sin letras y las invariables %s. Los literales de '
                                             'fórmula están en censo_es.json (mapas.CLAVES).' % sorted(INVARIABLES))])),
        ('campos', OrderedDict([
            ('donde', 'f = libro, h = hoja, c = celda / sqref / campo; t = celda | hoja | dv-* | cabecera-pie | '
                      'docprops; tr = tratamiento de esa aparición; det = detalle'),
            ('en_mapa', 'valor EN obligatorio que dicta mapas.py'),
            ('cita_hojas', 'pestañas ES citadas entre comillas, con su EN'),
            ('pistas', 'decisiones de la SPEC que aplican a la cadena'),
        ])),
    ])
    textos['cadenas'] = cadenas

    os.makedirs(args.salida, exist_ok=True)
    for nombre, obj in (('textos_es.json', textos), ('censo_es.json', censo)):
        with open(os.path.join(args.salida, nombre), 'w', encoding='utf-8') as fh:
            json.dump(obj, fh, ensure_ascii=False, indent=1)
            fh.write('\n')

    print('textos_es.json: {} cadenas ({} traducir, {} localizar, {} palabras para subagentes) · {} apariciones'.format(
        len(cadenas), textos['meta']['n_traducir'], textos['meta']['n_localizar'],
        textos['meta']['n_palabras_subagentes'], textos['meta']['n_apariciones']))
    for g in grupos:
        print('  {:<3} {:>4} cadenas {:>6} palabras  {}'.format(g['id'], g['n_cadenas'], g['n_palabras'],
                                                              g['descripcion'][:110]))
    print('censo_es.json: ' + json.dumps({k: v for k, v in T.items() if not isinstance(v, dict)},
                                         ensure_ascii=False))
    if args.md:
        print(tabla_md(censo))


COLS_MD = [('hojas', 'Hojas'), ('celdas_texto', 'Texto'), ('celdas_formula', 'Fórmulas'),
           ('formulas_con_hoja', 'Fx→hoja'), ('dv', 'DV'), ('cf', 'CF'), ('graficos', 'Gráficos'), ('merges', 'Merges'),
           ('hojas_protegidas', 'Protegidas'), ('fechas', 'Fechas'),
           ('formatos_euro', 'Fmt €'), ('textos_moneda', 'Textos €'), ('textos_normativa_es', 'Norma ES'),
           ('verdes', 'Verdes'), ('desbloqueadas', 'Desbloq.')]


def tabla_md(censo):
    out = ['| Libro | ' + ' | '.join(t for _, t in COLS_MD) + ' |', '|---|' + '---|' * len(COLS_MD)]
    for f, info in censo['libros'].items():
        out.append('| %s | ' % f[:-5] + ' | '.join(str(info['totales'][k]) for k, _ in COLS_MD) + ' |')
    out.append('| **Total** | ' + ' | '.join('**%s**' % censo['totales'][k] for k, _ in COLS_MD) + ' |')
    return '\n'.join(out)


if __name__ == '__main__':
    main()
