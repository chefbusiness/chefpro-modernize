#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
gates_en.py — Gates F2 del Restaurant Financial Plan Kit Pro (EN), SPEC §8 (G1-G8), sobre la salida de aplicar_en.py.
Copia adaptada de `staff-kit/gates_en.py` (+ gráficos y TIR). Solo LEE: los 10 xlsx EN, los 10 xlsx ES publicados
(dl/kit-plan-financiero, solo lectura), censo_es.json y mapas.py. Ningún recuento está escrito a mano: todo sale de
censo_es.json, de los xlsx ES o de mapas.py; las únicas cifras propias son las de SPEC §5 y §8 G6 (datos de ejemplo,
caso de la TIR, caso D10, 06 C30 = ventas / 860).

    python3 gates_en.py                               # contra astro-site/public/dl/restaurant-financial-plan-templates/
    python3 gates_en.py --dir /tmp/fin-dry            # contra un dry-run
    python3 gates_en.py --solo G1,G2 [--mes October] [--json informe.json]
    python3 gates_en.py --autotest [--dir …]          # inyecta defectos en una COPIA y exige que se detecten

Gates (SPEC §8):
  G1 paridad: 56 hojas por §2.2; fórmula a fórmula = ES tras HOJAS + CLAVES salvo el patrón D10 (48 celdas);
     fórmulas con hoja; DV: las 60 del ES (tipo, flags, fórmulas por CLAVES/HOJAS, mensajes presentes) con las 5
     DV_PARTIDAS (C conserva la DV 0-1, H estrena la de vida útil 0-50; unión de sqref = la del ES); CF (sqref,
     tipo, prioridad, relleno, fórmula y `text`); merges; títulos de impresión; paneles; columnas ocultas;
     protección (56 hojas, sin contraseña, mismos flags); verdes y desbloqueadas; 9 gráficos y 20 series (hoja,
     tipo, estilo, ancla, títulos de eje, referencias con pestaña EN); 0 imágenes;
  G2 restos: cero español, «€», «m²», «p.p.», caracteres fuera de la lista blanca, coma decimal, normativa ES
     (RX_NORMA) e IDs internos en valores, literales, DV (ítem a ítem), CF, gráficos, formatos, pies y docProps;
     celdas con formato € reescritas; citas a otro fichero de UN nivel; DV con título ≤ 32 y mensaje ≤ 255;
     DV_MENSAJES_EN;
  G3 formato: paperSize 1 en las 56; cada formato ES de FORMATOS_EN con su EN celda a celda (fechas numFmtId 14,
     «p.p.» → 0.0 % en las 64 de I1, « años» → « years»); 04 H sin %; ningún otro formato de fecha; pie D18;
     marca D20 donde el ES la tenía; versión D20; docProps D5/D20; Instructions!B2 = título interior D5;
     AJUSTE_TEXTO con ajuste y alto de fila;
  G4 claves: ítems de la DV del B09 ∈ CLAVES; cada estado escrito = un ítem; cada número de una DV numérica dentro
     de sus límites (vida útil 0-50 incluida);
  G5 CF: tokens EN inyectivos por rango, todos ∈ CLAVES (o sin letras) y mapas.cruzar_censo;
  G6 cálculo: <v> en toda <f>, cero errores; caché EN = caché de REFERENCIA (el ES publicado con VALORES_EN y el
     patrón D10 aplicados, recalculado con pycel) celda a celda; cifras de SPEC §5 (31,460 · break-even · EBITDA
     17.3 % · saldo final 15,000 · simulador · 06 C30 = ventas / 860); TIR: caso trazado del ES (−150,000 / 30,000 /
     45,000 / 60,000 / 70,000 → 11.9592 %) en copias del 07 EN y ES (misma TIR, = Newton y = bisección
     independiente; VAN y payback = Python; DSCR con el préstamo SBA de ejemplo); caso D10 (7,000 / 7 años → 1,000;
     vida vacía → 0);
  G7 reglas: VALORES_EN escritos y cifras = SPEC §3 (mapas.REGLAS_US); patrón D10 en las 48 celdas; DV de vida útil;
  G8 censo-entregables.py --fail (Letter) y gate-no-latinos.py en 0.
Sale con código 1 si algún gate falla.
"""
import argparse
import datetime
import html
import json
import logging
import os
import re
import shutil
import subprocess
import sys
import tempfile
import unicodedata
import zipfile
from collections import OrderedDict, Counter

AQUI = os.path.dirname(os.path.abspath(__file__))
SCRIPTS = os.path.dirname(AQUI)
REPO = os.path.dirname(os.path.dirname(SCRIPTS))
sys.path.insert(0, AQUI)

import openpyxl                                          # noqa: E402

import mapas                                             # noqa: E402
import aplicar_en as A                                   # noqa: E402

logging.disable(logging.CRITICAL)

ORIGEN = A.ORIGEN
CENSO_SCRIPT = os.path.join(SCRIPTS, 'censo-entregables.py')
NO_LATINOS = os.path.join(SCRIPTS, 'gate-no-latinos.py')
TOL = 0.006
ERRORES = ('#REF!', '#VALUE!', '#NAME?', '#DIV/0!', '#N/A', '#NUM!', '#NULL!')
CLAVES = mapas.CLAVES
CLAVES_EN = set(CLAVES.values())

SIGNOS_ES = set('¿¡«»€ºª')
RX_ID_INTERNO = re.compile(r'\[(?:fuente|estimado|derivado|criterio|kit estimate)[^\]]*\]|\bSPEC\b|\((?:D|I)\d{1,2}\)|'
                           r'\bD\d{1,2}\b(?=[:,)])|\bTEC-\d+\b|\bkit estimate\b|'
                           r'\b(?:kitpf|kitgp|grupo[A-Z])-v\d+[a-z]?\b|\bRT-\d+\b|\bI\d\b(?=[:,)])')
# Citas a otro fichero: file NN "nombre", UN nivel (staff-kit §6.5)
RX_CITA_ANIDADA = re.compile(r'\(\(|\)\)\)|\bfile (?:0\d[b]?|BONUS-0\d) \(')
LIMITE_DV = {'promptTitle': 32, 'errorTitle': 32, 'prompt': 255, 'error': 255}     # Excel
RX_TILDE = re.compile(r'[áéíóúüñÁÉÍÓÚÜÑ]')
LISTA_BLANCA_TILDE = {'café', 'cafés', 'entrée', 'entrées', 'sauté', 'sautéed', 'purée', 'jalapeño', 'jalapeños',
                      'crème', 'brûlée', 'résumé', 'résumés', 'naïve'}
FUNCIONALES_ES = {
    'de', 'del', 'las', 'el', 'los', 'y', 'que', 'con', 'para', 'por', 'una', 'unos', 'unas', 'es', 'al', 'se',
    'su', 'sus', 'como', 'pero', 'muy', 'cuando', 'donde', 'porque', 'este', 'esta', 'esto', 'estos', 'desde',
    'hasta', 'sobre', 'entre', 'tu', 'tus', 'te', 'mis', 'hoja', 'pestaña', 'celda', 'celdas', 'sin', 'ejemplo',
    'cada', 'debe', 'hay', 'puede', 'si', 'lo', 'le', 'les', 'todo', 'toda', 'todos', 'nunca', 'siempre',
}
OFICIO_ES = {
    'ventas', 'venta', 'facturacion', 'gastos', 'gasto', 'ingresos', 'ingreso', 'coste', 'costes', 'prestamo',
    'inversion', 'tesoreria', 'amortizacion', 'beneficio', 'beneficios', 'impuesto', 'impuestos', 'cuota', 'saldo',
    'cobros', 'cobro', 'pagos', 'pago', 'proveedores', 'proveedor', 'alquiler', 'suministros', 'nominas', 'nomina',
    'plantilla', 'plantillas', 'cubiertos', 'comedor', 'barra', 'mes', 'meses', 'anual', 'mensual', 'resumen',
    'datos', 'escenarios', 'escenario', 'parametros', 'flujo', 'alertas', 'obra', 'licencias', 'licencia',
    'equipamiento', 'mobiliario', 'sala', 'tecnologia', 'fondos', 'propios', 'deuda', 'banco', 'bancos', 'garantias',
    'simulador', 'comparativa', 'pendiente', 'completada', 'curso', 'presupuesto', 'desviacion', 'dotacion',
    'partida', 'concepto', 'tasa', 'tipo', 'plazo', 'cuadro', 'intereses', 'objetivo', 'limite', 'sentido', 'estado',
    'fase', 'tarea', 'tareas', 'responsable', 'notas', 'fecha', 'gestoria', 'socio', 'socios', 'promotor', 'aforo',
    'plazas', 'superficie', 'ubicacion', 'apertura', 'constitucion', 'seguro', 'seguros', 'contrato', 'notario',
    'registro', 'calendario', 'liquidacion', 'retenciones', 'tarjeta', 'efectivo', 'compras', 'existencias',
    'consumo', 'bebida', 'comida', 'otros', 'previsto', 'prevista', 'previstos', 'previstas', 'umbral', 'minimo',
    'bajo', 'alerta', 'aviso', 'semaforo', 'dias', 'dia', 'horas', 'medio', 'margen', 'contribucion', 'punto',
    'equilibrio', 'calculadora', 'liquidez', 'endeudamiento', 'solvencia', 'rentabilidad', 'proyecciones',
    'financiacion', 'garantia', 'ejecutivo', 'informe', 'viabilidad', 'cocina', 'hosteleria', 'negocio', 'mesa',
    'materia', 'prima', 'mermas', 'carencia', 'traspaso', 'fianza', 'aval',
}
RX_URL = re.compile(r'https?://\S+|\b[\w-]+(?:\.[\w-]+)*\.(?:es|com|pro|gov|org|uk|co|example|net)\b[/\w.#-]*', re.I)
RX_PALABRA = re.compile(r"[A-Za-zÀ-ÖØ-öø-ÿ]+")
RX_STR = A.RX_STR
RX_NORMA = re.compile(A.extraer_textos.RX_NORMA.pattern)
RX_SIGLAS_ES = re.compile(r'\b(?:TPV|IVA|CIF|NIF|IRPF|TIN|TAE|SGR|ICO|RETA|CCC|BAI)\b')
RX_COMA_DECIMAL = re.compile(r'(?<![\d,])\d+,\d{1,2}(?![\d])')
RX_PP = re.compile(r'\bp\.\s?p\.')


def _simbolos_es():
    """Lista blanca de caracteres no ASCII que no son letras: los del ES publicado (▸ → ✓ ✗ ⚠ emojis…) +
    la tipografía del kit EN (§ ≤ ≥ − ± × ’ “ ”). Los signos españoles (¿ ¡ « » € º ª) nunca entran."""
    ok = set('▸→☐✓✗⚠·—–…°§≤≥−±×½✅❌↑↓⭐') | {'️'}
    te = json.load(open(os.path.join(AQUI, 'textos_es.json'), encoding='utf-8'))
    textos = [c['es'] for c in te['cadenas']] + list(te['meta']['excluidas']['cadenas'])
    censo = json.load(open(os.path.join(AQUI, 'censo_es.json'), encoding='utf-8'))
    for info in censo['libros'].values():
        for hd in info['hojas_detalle'].values():
            textos += list(hd['literales']) + list(hd['literales_cf'])
    for t in textos:
        for ch in t:
            if ord(ch) > 127 and not unicodedata.category(ch).startswith('L') and ch not in SIGNOS_ES:
                ok.add(ch)
    ok.discard('²')                                   # m² (D14) no vuelve en la EN
    return ok


SIMBOLOS_OK = _simbolos_es()


def no_permitidos(t):
    return sorted({ch for ch in t if ord(ch) > 127 and ch not in SIMBOLOS_OK
                   and not ('À' <= ch <= 'ÿ' and ch not in '×÷')})


def restos_espanol(t):
    out = []
    limpio = RX_URL.sub(' ', t)
    for ch in limpio:
        if ch in SIGNOS_ES:
            out.append(('signo «%s»' % ch, t[:80]))
    for m in RX_PALABRA.finditer(limpio):
        w = m.group(0)
        wl = w.lower()
        ctx = limpio[max(0, m.start() - 30):m.end() + 30].replace('\n', ' ')
        sin = (wl.replace('á', 'a').replace('é', 'e').replace('í', 'i').replace('ó', 'o').replace('ú', 'u')
               .replace('ñ', 'n'))
        if RX_TILDE.search(w) and wl not in LISTA_BLANCA_TILDE:
            out.append(('tilde/eñe «%s»' % w, ctx))
        elif wl in FUNCIONALES_ES and w not in ('Y', 'T', 'N', 'NO', 'SE', 'TE', 'SI', 'S'):
            out.append(('funcional ES «%s»' % w, ctx))
        elif sin in OFICIO_ES and wl not in LISTA_BLANCA_TILDE:
            out.append(('oficio ES «%s»' % w, ctx))
    m = RX_COMA_DECIMAL.search(t)
    if m:
        out.append(('coma decimal «%s»' % m.group(0), t[:80]))
    m = RX_SIGLAS_ES.search(t)
    if m:
        out.append(('sigla ES «%s»' % m.group(0), t[:80]))
    m = RX_NORMA.search(t)
    if m:
        out.append(('normativa ES «%s»' % m.group(0), t[:80]))
    if 'm²' in t:
        out.append(('m² (D14: sq ft)', t[:80]))
    if RX_PP.search(t):
        out.append(('«p.p.» (D18: pts)', t[:80]))
    return out


# ==========================================================================
# Utilidades
# ==========================================================================
class Resultado:
    def __init__(self):
        self.gates = OrderedDict()
        self.actual = None

    def gate(self, nombre):
        self.actual = nombre
        self.gates.setdefault(nombre, {'fallos': [], 'datos': OrderedDict()})
        return self

    def fallo(self, msg):
        self.gates[self.actual]['fallos'].append(msg)

    def dato(self, k, v):
        self.gates[self.actual]['datos'][k] = v


_CACHE = {}


def cargar(path, data_only=False):
    k = (path, data_only, os.path.getmtime(path))
    if k not in _CACHE:
        _CACHE[k] = openpyxl.load_workbook(path, data_only=data_only)
    return _CACHE[k]


def libros(carpeta):
    for es in mapas.LIBROS_ES:
        en = mapas.FICHEROS[es]
        yield es, en, A.extraer_textos.corto(es), os.path.join(ORIGEN, es), os.path.join(carpeta, en)


def fichero_en(carpeta, corto):
    return [pen for es, en, c, pes, pen in libros(carpeta) if c == corto][0]


def fichero_es(corto):
    return [pes for es, en, c, pes, pen in libros('.') if c == corto][0]


def lits(f):
    return [x[1:-1].replace('""', '"') for x in RX_STR.findall(f or '')]


def celdas(sqref):
    return A.rango_lista(sqref)


def es_verde(c):
    return (c.fill is not None and c.fill.fgColor is not None and isinstance(c.fill.fgColor.rgb, str)
            and c.fill.fgColor.rgb.upper().endswith(A.VERDE))


def celdas_reales(ws):
    for row in ws.iter_rows():
        for c in row:
            if c.__class__.__name__ != 'MergedCell':
                yield c


def _pagina(ws):
    ps = ws.page_setup
    fit = ws.sheet_properties.pageSetUpPr.fitToPage if ws.sheet_properties.pageSetUpPr else None
    return (ps.orientation, ps.fitToWidth, ps.fitToHeight, fit, str(ws.print_title_rows), str(ws.print_title_cols))


def _area(ws):
    pa = ws.print_area
    return pa.split('!')[-1] if pa else None


def _proteccion(ws):
    p = ws.protection
    return (bool(p.sheet), bool(p.password), p.formatColumns, p.formatRows, p.insertRows, p.deleteRows, p.sort,
            p.autoFilter, p.selectLockedCells, p.selectUnlockedCells)


def _cf_firma(cf, regla, trans, trans_text):
    dxf = regla.dxf
    color = None
    if dxf is not None:
        color = (str(dxf.fill.fgColor.rgb) if dxf.fill is not None and dxf.fill.fgColor is not None else None,
                 str(dxf.font.color.rgb) if dxf.font is not None and dxf.font.color is not None else None)
    return (str(cf.sqref), regla.type, regla.operator, regla.priority, bool(regla.stopIfTrue),
            tuple(trans(x) for x in (regla.formula or [])), trans_text(regla.text) if regla.text else regla.text, color)


def dv_firma(d, sqref=None, tipo=None, f=None):
    """Firma de una DV: sqref, flags, tipo y fórmulas (las del ES transformadas o las del EN tal cual)."""
    t, op, f1, f2 = f if f is not None else (d.type, d.operator, d.formula1, d.formula2)
    return (sqref if sqref is not None else str(d.sqref), bool(d.allow_blank), bool(d.showDropDown),
            bool(d.showErrorMessage), bool(d.showInputMessage), d.errorStyle, tipo or t, op, f1, f2,
            bool(d.error), bool(d.errorTitle), bool(d.prompt), bool(d.promptTitle))


def dv_en_esperada(de, T):
    """(type, operator, formula1, formula2) EN de una DV ES (misma lógica que aplicar_en, recalculada aquí)."""
    f1 = de.formula1 or ''
    if de.type == 'list' and f1.startswith('"'):
        items = [CLAVES[x] if mapas.necesita_clave(x) else x for x in f1.strip('"').split(',')]
        return de.type, de.operator, mapas.dv_lista(items), de.formula2
    if de.type == 'custom' and f1:
        return de.type, de.operator, T.formula(f1), de.formula2
    return (de.type, de.operator, T.refs(f1) if f1 else de.formula1, T.refs(de.formula2) if de.formula2 else de.formula2)


def serial(v):
    if isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        return v
    if isinstance(v, datetime.datetime):
        return (v - datetime.datetime(1899, 12, 30)).total_seconds() / 86400.0
    if isinstance(v, datetime.date):
        return (v - datetime.date(1899, 12, 30)).days
    return None


def num(v):
    return isinstance(v, (int, float)) and not isinstance(v, bool)


# ---- gráficos: ancla y estilo leídos del ZIP
def _xml(z, n):
    return z.read(n).decode('utf-8')


def anclas_graficos(path):
    """{hoja: [ancla normalizada]} — from/to/ext de cada objeto del drawing de la hoja."""
    out = {}
    with zipfile.ZipFile(path) as z:
        nombres = set(z.namelist())
        wbx, rels = _xml(z, 'xl/workbook.xml'), _xml(z, 'xl/_rels/workbook.xml.rels')
        rid2t = {A.extraer_textos._xml_attr(t, 'Id'): A.extraer_textos._xml_attr(t, 'Target')
                 for t in re.findall(r'<Relationship\b[^>]*>', rels)}
        for t in re.findall(r'<sheet\b[^>]*>', wbx):
            hoja = html.unescape(A.extraer_textos._xml_attr(t, 'name'))
            parte = A.extraer_textos._resolver('xl/workbook.xml', rid2t[A.extraer_textos._xml_attr(t, 'r:id')])
            srel = os.path.dirname(parte) + '/_rels/' + os.path.basename(parte) + '.rels'
            if srel not in nombres:
                continue
            for r in re.findall(r'<Relationship\b[^>]*>', _xml(z, srel)):
                if A.extraer_textos._xml_attr(r, 'Type').endswith('/drawing'):
                    d = _xml(z, A.extraer_textos._resolver(parte, A.extraer_textos._xml_attr(r, 'Target')))
                    for m in re.finditer(r'<(?:xdr:)?(oneCellAnchor|twoCellAnchor|absoluteAnchor)\b[^>]*>(.*?)'
                                         r'</(?:xdr:)?\1>', d, re.S):
                        trozos = re.findall(r'<(?:xdr:)?(?:from|to)>.*?</(?:xdr:)?(?:from|to)>|<(?:xdr:)?ext\b[^>]*/>',
                                            m.group(2), re.S)
                        out.setdefault(hoja, []).append(m.group(1) + ':' + re.sub(r'xdr:', '', ''.join(trozos)))
    return out


def estilos_graficos(path):
    with zipfile.ZipFile(path) as z:
        return {n: (re.findall(r'<(?:c:)?style val="(\d+)"', _xml(z, n)), re.findall(r'<(?:c:)?(\w+Chart)>', _xml(z, n)),
                    len(re.findall(r'<(?:c:)?ser>', _xml(z, n))))
                for n in z.namelist() if re.fullmatch(r'xl/charts/chart\d+\.xml', n)}


def textos_grafico(path):
    with zipfile.ZipFile(path) as z:
        for n in z.namelist():
            if re.fullmatch(r'xl/charts/chart\d+\.xml', n):
                x = _xml(z, n)
                for t in re.findall(r'<a:t>([^<]*)</a:t>', x):
                    yield n, 'texto', html.unescape(t)
                for t in re.findall(r'<(?:c:)?v>([^<]*)</(?:c:)?v>', x):
                    yield n, 'cache', html.unescape(t)
                for t in re.findall(r'formatCode="([^"]*)"', x):
                    yield n, 'formato', html.unescape(t)


# ==========================================================================
# G1 — paridad estructural
# ==========================================================================
def gateG1(R, carpeta, datos):
    R.gate('G1 paridad')
    censo = datos['censo']
    tot = Counter()
    for es, en, corto, pes, pen in libros(carpeta):
        if not os.path.isfile(pen):
            R.fallo('%s: no existe' % en)
            continue
        ce = censo['libros'][es]
        wes, wen = cargar(pes), cargar(pen)
        esperado = [mapas.HOJAS[h] for h in ce['hojas']]
        if wen.sheetnames != esperado or wes.sheetnames != ce['hojas']:
            R.fallo('%s: hojas %r ≠ %r' % (en, wen.sheetnames, esperado))
            continue
        tot['hojas'] += len(esperado)
        mapa = OrderedDict((h, mapas.HOJAS[h]) for h in wes.sheetnames)
        T = A.Transformador(mapa)
        for h_es in wes.sheetnames:
            ws_es, ws_en = wes[h_es], wen[mapa[h_es]]
            hd = ce['hojas_detalle'][h_es]
            lugar = '%s:%s' % (en, ws_en.title)
            f_es = {c.coordinate: c.value for c in ws_es._cells.values() if c.data_type == 'f'}
            f_en = {c.coordinate: c.value for c in ws_en._cells.values() if c.data_type == 'f'}
            if len(f_es) != hd['celdas_formula']:
                R.fallo('%s: el ES ya no cuadra con el censo (%d ≠ %d fórmulas)' % (lugar, len(f_es), hd['celdas_formula']))
            if set(f_en) != set(f_es):
                R.fallo('%s: celdas con fórmula distintas: solo ES %s · solo EN %s'
                        % (lugar, sorted(set(f_es) - set(f_en))[:5], sorted(set(f_en) - set(f_es))[:5]))
            con_hoja = 0
            for k, fe in f_es.items():
                fn = f_en.get(k)
                if fn is None:
                    continue
                tot['formulas'] += 1
                try:
                    esperado_f, es_pat = A.formula_en(fe, ws_es[k].row, T)
                except A.Aborta as e:
                    esperado_f, es_pat = '<<%s>>' % e, False
                tot['formulas_patron_en'] += es_pat
                if esperado_f != fn:
                    R.fallo('%s!%s: fórmula EN %r ≠ esperada %r' % (lugar, k, fn[:80], esperado_f[:80]))
                if A.extraer_textos.refs_hoja(fn):
                    con_hoja += 1
                    for h in A.extraer_textos.refs_hoja(fn):
                        if h not in wen.sheetnames:
                            R.fallo('%s!%s cita la hoja inexistente %r' % (lugar, k, h))
            tot['formulas_con_hoja'] += con_hoja
            if con_hoja != hd['formulas_con_hoja']:
                R.fallo('%s: %d fórmulas con referencia a hoja ≠ censo %d' % (lugar, con_hoja, hd['formulas_con_hoja']))
            # DV: las del ES transformadas; las DV_PARTIDAS se parten en C (misma DV) + H (vida útil)
            esperadas = []
            for de in ws_es.data_validations.dataValidation:
                sq = str(de.sqref)
                try:
                    f = dv_en_esperada(de, T)
                except (KeyError, A.Aborta) as e:
                    f = ('<<%s>>' % e, None, None, None)
                partida = mapas.DV_PARTIDAS.get((corto, h_es, sq))
                if partida:
                    esperadas.append(dv_firma(de, sqref=partida[0], f=f))
                    esperadas.append(dv_firma(de, sqref=partida[1], f=partida[2:6]) [:10] + (True, True, False, False))
                    tot['dv_partidas'] += 1
                else:
                    esperadas.append(dv_firma(de, f=f))
            reales = [dv_firma(d) for d in ws_en.data_validations.dataValidation]
            if sorted(map(repr, esperadas)) != sorted(map(repr, reales)):
                faltan = sorted(set(map(repr, esperadas)) - set(map(repr, reales)))
                sobran = sorted(set(map(repr, reales)) - set(map(repr, esperadas)))
                R.fallo('%s: DV distintas · esperada %s · real %s' % (lugar, faltan[:1], sobran[:1]))
            if len(ws_es.data_validations.dataValidation) != len(hd['dv']):
                R.fallo('%s: el ES ya no cuadra con el censo en DV' % lugar)
            union_es = set(c for d in ws_es.data_validations.dataValidation for c in celdas(d.sqref))
            union_en = set(c for d in ws_en.data_validations.dataValidation for c in celdas(d.sqref))
            if union_es != union_en:
                R.fallo('%s: la unión de sqref de las DV cambia (%s)' % (lugar, sorted(union_es ^ union_en)[:5]))
            tot['dv'] += len(reales)
            tot['dv_celdas'] += sum(len(celdas(d.sqref)) for d in ws_en.data_validations.dataValidation)
            # CF
            cfe = sorted((_cf_firma(cf, r, T.formula, lambda t: A.clave_en(t, 'CF')) for cf in ws_es.conditional_formatting
                          for r in cf.rules), key=repr)
            cfn = sorted((_cf_firma(cf, r, lambda x: x, lambda t: t) for cf in ws_en.conditional_formatting
                          for r in cf.rules), key=repr)
            tot['cf'] += len(cfn)
            if len(cfn) != len(hd['cf']):
                R.fallo('%s: %d CF EN ≠ censo %d' % (lugar, len(cfn), len(hd['cf'])))
            if cfe != cfn:
                dif = [(a, b) for a, b in zip(cfe, cfn) if a != b] or [(cfe[:1], cfn[:1])]
                R.fallo('%s: formato condicional distinto (ES transformado → EN): %r' % (lugar, dif[0]))
            # merges, protección, paneles, impresión, columnas ocultas
            mg = sorted(map(str, ws_en.merged_cells.ranges))
            if mg != sorted(hd['merges']):
                R.fallo('%s: combinaciones distintas' % lugar)
            tot['merges'] += len(mg)
            if _proteccion(ws_en) != _proteccion(ws_es) or ws_en.protection.password \
                    or bool(ws_en.protection.sheet) != hd['proteccion']['hoja']:
                R.fallo('%s: protección distinta del ES %r ≠ %r' % (lugar, _proteccion(ws_en), _proteccion(ws_es)))
            tot['hojas_protegidas'] += 1 if ws_en.protection.sheet else 0
            if ws_en.freeze_panes != hd['paneles']:
                R.fallo('%s: paneles %s ≠ %s' % (lugar, ws_en.freeze_panes, hd['paneles']))
            tot['paneles'] += 1 if ws_en.freeze_panes else 0
            if _pagina(ws_es) != _pagina(ws_en):
                R.fallo('%s: ajuste o títulos de impresión %r ≠ %r' % (lugar, _pagina(ws_es), _pagina(ws_en)))
            tot['titulos_impresion'] += 1 if ws_en.print_title_rows else 0
            ae = hd['area_impresion'].split('!')[-1] if hd['area_impresion'] else None
            if _area(ws_en) != ae:
                R.fallo('%s: área de impresión %s ≠ %s' % (lugar, _area(ws_en), ae))
            tot['areas_impresion'] += 1 if ae else 0
            ocultas = sorted(k for k, d in ws_en.column_dimensions.items() if d.hidden)
            if ocultas != hd['columnas_ocultas']:
                R.fallo('%s: columnas ocultas %s ≠ %s' % (lugar, ocultas, hd['columnas_ocultas']))
            tot['columnas_ocultas'] += len(ocultas)
            ue = {c.coordinate for c in celdas_reales(ws_es) if c.protection.locked is False}
            un = {c.coordinate for c in celdas_reales(ws_en) if c.protection.locked is False}
            ve = {c.coordinate for c in celdas_reales(ws_es) if es_verde(c)}
            vn = {c.coordinate for c in celdas_reales(ws_en) if es_verde(c)}
            if un != ue:
                R.fallo('%s: desbloqueadas distintas (solo ES %s, solo EN %s)' % (lugar, sorted(ue - un)[:5],
                                                                                   sorted(un - ue)[:5]))
            if vn != ve:
                R.fallo('%s: verdes distintas (solo ES %s, solo EN %s)' % (lugar, sorted(ve - vn)[:5], sorted(vn - ve)[:5]))
            if vn - un:
                R.fallo('%s: celdas verdes bloqueadas %s' % (lugar, sorted(vn - un)[:5]))
            tot['verdes'] += len(vn)
            tot['desbloqueadas'] += len(un)
        # gráficos (D19): hoja, tipos, estilo, series, ejes, referencias con pestaña EN, ancla
        g_es, g_en = A.extraer_textos.graficos(pes), A.extraer_textos.graficos(pen)
        tot['graficos'] += len(g_en)
        tot['series_grafico'] += sum(g['series'] for g in g_en)
        if len(g_es) != len(g_en):
            R.fallo('%s: %d gráficos ≠ %d del ES' % (en, len(g_en), len(g_es)))
        st_es, st_en = estilos_graficos(pes), estilos_graficos(pen)
        por_hoja_en = {g['hoja']: g for g in g_en}
        for g in g_es:
            gn = por_hoja_en.get(mapa[g['hoja']])
            if gn is None:
                R.fallo('%s: el gráfico de «%s» no está en «%s»' % (en, g['hoja'], mapa[g['hoja']]))
                continue
            if (gn['tipos'], gn['series'], len(gn['titulos_eje'])) != (g['tipos'], g['series'], len(g['titulos_eje'])):
                R.fallo('%s: gráfico %s tipo/series/ejes %r ≠ ES %r' % (en, gn['parte'], (gn['tipos'], gn['series']),
                                                                       (g['tipos'], g['series'])))
            refs_esp = [T.formula(r) for r in g['refs']]
            if gn['refs'] != refs_esp:
                R.fallo('%s: gráfico %s referencias %r ≠ %r' % (en, gn['parte'], gn['refs'][:3], refs_esp[:3]))
            if not gn['titulo'] or (g['titulo'] and gn['titulo'] == g['titulo']):
                R.fallo('%s: gráfico %s sin título EN (%r)' % (en, gn['parte'], gn['titulo']))
            if st_en.get(gn['parte']) != st_es.get(g['parte']):
                R.fallo('%s: gráfico %s estilo/tipo %r ≠ ES %r' % (en, gn['parte'], st_en.get(gn['parte']),
                                                                    st_es.get(g['parte'])))
        an_es, an_en = anclas_graficos(pes), anclas_graficos(pen)
        if {mapa[h]: v for h, v in an_es.items()} != an_en:
            R.fallo('%s: anclas de gráfico distintas %r ≠ %r' % (en, an_en, an_es))
        with zipfile.ZipFile(pen) as z:
            imgs = [x for x in z.namelist() if x.startswith('xl/media/')]
        if imgs:
            R.fallo('%s: imágenes %s' % (en, imgs[:3]))
    T0 = censo['totales']
    n_pat = sum(p['n'] for info in censo['libros'].values() for hd in info['hojas_detalle'].values()
                for p in hd['patrones_formula'] if p['patron'] in mapas.FORMULAS_EN)
    esperados = OrderedDict([
        ('hojas', T0['hojas']), ('formulas', T0['celdas_formula']), ('formulas_patron_en', n_pat),
        ('formulas_con_hoja', T0['formulas_con_hoja']), ('dv', T0['dv'] + len(mapas.DV_PARTIDAS)),
        ('dv_partidas', len(mapas.DV_PARTIDAS)), ('dv_celdas', T0['dv_celdas']), ('cf', T0['cf']),
        ('merges', T0['merges']), ('hojas_protegidas', T0['hojas_protegidas']),
        ('areas_impresion', T0['areas_impresion']), ('titulos_impresion', T0['titulos_impresion']),
        ('paneles', T0['paneles']), ('columnas_ocultas', T0['columnas_ocultas']),
        ('verdes', T0['verdes']), ('desbloqueadas', T0['desbloqueadas']),
        ('graficos', T0['graficos']), ('series_grafico', T0['series_grafico']),
    ])
    for k, v in esperados.items():
        if tot[k] != v:
            R.fallo('total %s = %d ≠ esperado %d (censo)' % (k, tot[k], v))
    for k, v in sorted(tot.items()):
        R.dato(k, v)


# ==========================================================================
# G2 — restos y caracteres
# ==========================================================================
def piezas_de_texto(pen):
    wb = cargar(pen)
    for ws in wb.worksheets:
        yield 'hoja', ws.title, ws.title
        for c in ws._cells.values():
            if isinstance(c.value, str) and c.data_type != 'f':
                yield 'valor', '%s!%s' % (ws.title, c.coordinate), c.value
            if c.data_type == 'f':
                for x in lits(c.value):
                    yield 'literal', '%s!%s' % (ws.title, c.coordinate), x
            if isinstance(c.number_format, str) and c.number_format != 'General':
                yield 'formato', '%s!%s' % (ws.title, c.coordinate), c.number_format
        for dv in ws.data_validations.dataValidation:
            for attr in ('error', 'errorTitle', 'prompt', 'promptTitle'):
                v = getattr(dv, attr)
                if v:
                    yield 'dv', '%s DV %s %s' % (ws.title, dv.sqref, attr), v
            f1 = dv.formula1 or ''
            if f1.startswith('"'):
                for x in f1.strip('"').split(','):
                    if x:
                        yield 'dv-item', '%s DV %s' % (ws.title, dv.sqref), x
            elif dv.type == 'custom':
                for x in lits(f1):
                    yield 'literal', '%s DV %s' % (ws.title, dv.sqref), x
        for cf in ws.conditional_formatting:
            for regla in cf.rules:
                for f in regla.formula or []:
                    for x in lits(f):
                        yield 'literal', '%s CF %s' % (ws.title, cf.sqref), x
                if regla.text:
                    yield 'literal', '%s CF %s text' % (ws.title, cf.sqref), regla.text
        for parte in ('oddHeader', 'oddFooter', 'evenHeader', 'evenFooter', 'firstHeader', 'firstFooter'):
            hf = getattr(ws, parte)
            for pos in ('left', 'center', 'right'):
                t = getattr(hf, pos).text
                if t:
                    yield 'cabecera-pie', '%s %s.%s' % (ws.title, parte, pos), t
    p = wb.properties
    for k in ('title', 'subject', 'creator', 'keywords', 'description', 'category', 'lastModifiedBy'):
        v = getattr(p, k)
        if v:
            yield 'docprops', k, v
    for parte, tipo, t in textos_grafico(pen):
        yield ('formato' if tipo == 'formato' else 'grafico'), '%s %s' % (parte, tipo), t
    with zipfile.ZipFile(pen) as z:
        estilos = z.read('xl/styles.xml').decode('utf-8')
    for m in re.finditer(r'<numFmt [^>]*formatCode="([^"]*)"', estilos):
        yield 'styles.xml', 'numFmt', html.unescape(m.group(1))


def gateG2(R, carpeta, datos):
    R.gate('G2 restos')
    n = Counter()
    for es, en, corto, pes, pen in libros(carpeta):
        for tipo, lugar, t in piezas_de_texto(pen):
            n[tipo] += 1
            donde = '%s %s' % (en, lugar)
            if '€' in t or re.search(r'\bEUR\b|\beuros?\b', t):
                R.fallo('%s: moneda € en %s: %r' % (donde, tipo, t[:70]))
            malos = no_permitidos(t)
            if malos:
                R.fallo('%s: carácter fuera de la lista blanca %r en %s: %r' % (donde, ''.join(malos), tipo, t[:60]))
            if tipo in ('formato', 'styles.xml'):
                if re.search(r'dd/mm|yyyy-mm-dd|€|años|p\.p\.', t, re.I):
                    R.fallo('%s: formato %r' % (donde, t))
                continue
            m = RX_ID_INTERNO.search(t)
            if m:
                R.fallo('%s: ID interno %r en %s: …%s…' % (donde, m.group(0), tipo, t[max(0, m.start() - 30):m.end() + 10]))
            m = RX_CITA_ANIDADA.search(t)
            if m:
                R.fallo('%s: cita anidada %r en %s: …%s…' % (donde, m.group(0), tipo, t[max(0, m.start() - 40):m.end() + 30]))
            if tipo == 'dv':
                attr = lugar.rsplit(' ', 1)[-1]
                if len(t) > LIMITE_DV[attr]:
                    R.fallo('%s: DV %s de %d caracteres > %d (límite de Excel)' % (donde, attr, len(t), LIMITE_DV[attr]))
                n['dv_' + attr] += 1
            for motivo, ctx in restos_espanol(t):
                R.fallo('%s: %s · …%s…' % (donde, motivo, ctx))
        # D17: toda celda con formato € en el ES queda sin € en la EN
        wes, wen = cargar(pes), cargar(pen)
        for ws_es in wes.worksheets:
            ws_en = wen[mapas.HOJAS[ws_es.title]]
            for c in ws_es._cells.values():
                if c.number_format and '€' in c.number_format:
                    n['celdas_formato_euro_reescritas'] += 1
                    if '€' in (ws_en[c.coordinate].number_format or ''):
                        R.fallo('%s %s!%s: formato € %r' % (en, ws_en.title, c.coordinate, ws_en[c.coordinate].number_format))
    if n['celdas_formato_euro_reescritas'] != datos['censo']['totales']['formatos_euro']:
        R.fallo('celdas con formato € %d ≠ censo %d' % (n['celdas_formato_euro_reescritas'],
                                                        datos['censo']['totales']['formatos_euro']))
    for (corto, h_es, sq, attr), v in mapas.DV_MENSAJES_EN.items():
        ws = cargar(fichero_en(carpeta, corto))[mapas.HOJAS[h_es]]
        dvs = [d for d in ws.data_validations.dataValidation if str(d.sqref) == sq]
        n['dv_mensajes_fijados'] += 1
        if len(dvs) != 1 or getattr(dvs[0], attr) != v:
            R.fallo('DV_MENSAJES_EN %s %s DV %s %s = %r ≠ %r' % (corto, ws.title, sq, attr,
                                                                  getattr(dvs[0], attr) if dvs else None, v))
    for k, v in sorted(n.items()):
        R.dato(k, v)


# ==========================================================================
# G3 — formato
# ==========================================================================
def gateG3(R, carpeta, datos, mes):
    R.gate('G3 formato')
    censo = datos['censo']
    n = Counter()
    rx_version = re.compile(r'^Version %s · %s \d{4} · %s · info@aichef\.pro$'
                            % (re.escape(A.VERSION), mes, re.escape(A.DESCRIPTION)))
    por_rango = {}
    for (c, h, rg), fmt in mapas.FORMATOS_POR_RANGO.items():
        for coord in A.rango(rg):
            por_rango[(c, mapas.HOJAS[h], coord)] = fmt
    for es, en, corto, pes, pen in libros(carpeta):
        ce = censo['libros'][es]
        wb, wes = cargar(pen), cargar(pes)
        marcas_es = sum(1 for ws in wes.worksheets for c in ws._cells.values() if c.value == mapas.MARCA_ES)
        marcas_en = sum(1 for ws in wb.worksheets for c in ws._cells.values() if c.value == mapas.MARCA_EN)
        if marcas_en != marcas_es or marcas_es == 0:
            R.fallo('%s: %d líneas de marca D20 ≠ %d del ES' % (en, marcas_en, marcas_es))
        n['marcas'] += marcas_en
        for h_es, hd in ce['hojas_detalle'].items():
            ws, ws_es = wb[mapas.HOJAS[h_es]], wes[h_es]
            n['hojas'] += 1
            if ws.page_setup.paperSize not in (1, '1'):
                R.fallo('%s:%s: paperSize %s ≠ 1 (Letter)' % (en, ws.title, ws.page_setup.paperSize))
            fechas = {f['celda'] for f in hd['fechas']}
            for c_es in ws_es._cells.values():
                f_es = c_es.number_format
                c = ws[c_es.coordinate]
                k = (corto, ws.title, c_es.coordinate)
                if k in por_rango:
                    n['fmt_por_rango'] += 1
                    if c.number_format != por_rango[k]:
                        R.fallo('%s:%s!%s: formato %r ≠ %r (FORMATOS_POR_RANGO)' % (en, ws.title, c.coordinate,
                                                                                  c.number_format, por_rango[k]))
                    continue
                if f_es in A.FORMATOS:
                    esp = A.FORMATOS[f_es]
                    n['fmt_' + f_es] += 1
                    if esp in A.NUMFMT:
                        if c._style.numFmtId != A.NUMFMT[esp]:
                            R.fallo('%s:%s!%s: fecha con numFmtId %s ≠ %s' % (en, ws.title, c.coordinate,
                                                                              c._style.numFmtId, A.NUMFMT[esp]))
                    elif c.number_format != esp:
                        R.fallo('%s:%s!%s: formato %r ≠ %r (ES %r)' % (en, ws.title, c.coordinate, c.number_format,
                                                                       esp, f_es))
                elif c.number_format != f_es:
                    R.fallo('%s:%s!%s: formato %r cambia (ES %r)' % (en, ws.title, c.coordinate, c.number_format, f_es))
            for c in ws._cells.values():
                f = c.number_format or ''
                es_fecha = c._style.numFmtId in (14, 22) or re.search(r'[dmy]/[dmy]|h:mm', f)
                if es_fecha and c.coordinate not in fechas:
                    R.fallo('%s:%s!%s: formato de fecha %r en una celda que no lo era' % (en, ws.title, c.coordinate, f))
            for k in hd['cabeceras_pies']:
                obj, pos = k.split('.')
                v = getattr(getattr(ws, obj), pos).text
                if v != A.PIE_EN:
                    R.fallo('%s:%s: %s = %r ≠ pie D18' % (en, ws.title, k, v))
                else:
                    n['pies'] += 1
        ins = wb[A.HOJA_INSTR_EN]
        vers = [c.value for c in ins._cells.values() if isinstance(c.value, str) and c.value.startswith('Version ')]
        if len(vers) != 1 or not rx_version.match(vers[0]):
            R.fallo('%s: línea de versión D20 %r' % (en, vers))
        else:
            n['versiones'] += 1
        p = wb.properties
        esp = {'title': A.titulo_docprops(es), 'subject': A.SUBJECT, 'keywords': A.KEYWORDS,
               'description': A.DESCRIPTION, 'category': A.CATEGORY, 'creator': A.CREATOR}
        for k, v in esp.items():
            if getattr(p, k) != v:
                R.fallo('%s: docProps %s %r ≠ %r' % (en, k, getattr(p, k), v))
        if ins['B2'].value != mapas.titulo_interior(es):
            R.fallo('%s: Instructions!B2 %r ≠ título interior D5 %r' % (en, ins['B2'].value, mapas.titulo_interior(es)))
    for (corto, h_es, rg) in mapas.AJUSTE_TEXTO:
        ws = cargar(fichero_en(carpeta, corto))[mapas.HOJAS[h_es]]
        for coord in A.rango(rg):
            c = ws[coord]
            if c.value is None:
                continue
            n['ajuste_texto'] += 1
            alto = A.alto_necesario(ws, c.row)
            if not (c.alignment and c.alignment.wrap_text) or (alto > 15 and (ws.row_dimensions[c.row].height or 0) < alto):
                R.fallo('%s %s!%s: sin ajuste de texto o fila baja (alto %s < %s)' % (corto, ws.title, coord,
                                                                                   ws.row_dimensions[c.row].height, alto))
    for (corto, h_es), cols in mapas.ANCHOS.items():
        ws = cargar(fichero_en(carpeta, corto))[mapas.HOJAS[h_es]]
        for col, w in cols.items():
            n['anchos'] += 1
            if ws.column_dimensions[col].width != w:
                R.fallo('%s %s: columna %s de ancho %s ≠ ANCHOS %s' % (corto, ws.title, col,
                                                                      ws.column_dimensions[col].width, w))
    T0 = censo['totales']
    n_pies = sum(len(hd['cabeceras_pies']) for i in censo['libros'].values() for hd in i['hojas_detalle'].values())
    if n['hojas'] != T0['hojas'] or n['pies'] != n_pies:
        R.fallo('hojas %d / pies %d ≠ censo %d / %d' % (n['hojas'], n['pies'], T0['hojas'], n_pies))
    if n['versiones'] != len(mapas.FICHEROS):
        R.fallo('versiones %d ≠ %d' % (n['versiones'], len(mapas.FICHEROS)))
    n_pp = sum(v for i in censo['libros'].values() for hd in i['hojas_detalle'].values()
               for f, v in hd['formatos'].items() if f == '0.0" p.p."')
    if n['fmt_0.0" p.p."'] != n_pp or n['fmt_por_rango'] != sum(len(list(A.rango(r))) for (_c, _h, r) in mapas.FORMATOS_POR_RANGO):
        R.fallo('formatos: p.p. %d ≠ censo %d o FORMATOS_POR_RANGO incompleto' % (n['fmt_0.0" p.p."'], n_pp))
    if n['fmt_dd/mm/yyyy'] + n['fmt_DD/MM/YYYY'] != T0['fechas']:
        R.fallo('fechas %d ≠ censo %d' % (n['fmt_dd/mm/yyyy'] + n['fmt_DD/MM/YYYY'], T0['fechas']))
    for k, v in sorted(n.items()):
        R.dato(k, v)


# ==========================================================================
# G4 — claves
# ==========================================================================
def gateG4(R, carpeta, datos):
    R.gate('G4 claves')
    n = Counter()
    for es, en, corto, pes, pen in libros(carpeta):
        wb = cargar(pen)
        for ws in wb.worksheets:
            lugar = '%s:%s' % (en, ws.title)
            for dv in ws.data_validations.dataValidation:
                f1 = dv.formula1 or ''
                if dv.type == 'list' and f1.startswith('"'):
                    items = f1.strip('"').split(',')
                    n['dv_listas'] += 1
                    for x in items:
                        if mapas.necesita_clave(x) and x not in CLAVES_EN:
                            R.fallo('%s: DV %s ítem %r ∉ CLAVES' % (lugar, dv.sqref, x))
                    for coord in celdas(dv.sqref):
                        c = ws[coord]
                        if c.value is None or c.data_type == 'f':
                            continue
                        n['valores_con_dv_lista'] += 1
                        if str(c.value) not in items:
                            R.fallo('%s!%s: %r fuera de su DV %s' % (lugar, coord, c.value, f1[:60]))
                elif dv.type in ('decimal', 'whole') and f1:
                    try:
                        lo = float(f1)
                        hi = float(dv.formula2) if dv.formula2 else None
                    except ValueError:
                        continue
                    for coord in celdas(dv.sqref):
                        v = ws[coord].value
                        if v is None or ws[coord].data_type == 'f':
                            continue
                        n['valores_con_dv_numerica'] += 1
                        s = serial(v)
                        op = dv.operator or 'between'
                        ok = s is not None and (
                            (op == 'between' and lo - 1e-9 <= s <= hi + 1e-9) or
                            (op == 'greaterThanOrEqual' and s >= lo - 1e-9) or
                            (op == 'greaterThan' and s > lo) or (op == 'lessThanOrEqual' and s <= lo + 1e-9) or
                            op not in ('between', 'greaterThanOrEqual', 'greaterThan', 'lessThanOrEqual'))
                        if not ok:
                            R.fallo('%s!%s: %r fuera de su DV %s %s-%s' % (lugar, coord, v, op, f1, dv.formula2))
    for k, v in sorted(n.items()):
        R.dato(k, v)


# ==========================================================================
# G5 — CF
# ==========================================================================
def gateG5(R, carpeta, datos):
    R.gate('G5 CF')
    n = Counter()
    for es, en, corto, pes, pen in libros(carpeta):
        wes, wen = cargar(pes), cargar(pen)
        for h_es in wes.sheetnames:
            ws_es, ws_en = wes[h_es], wen[mapas.HOJAS[h_es]]
            tok_es, tok_en = {}, {}
            for ws, tok in ((ws_es, tok_es), (ws_en, tok_en)):
                for cf in ws.conditional_formatting:
                    for r in cf.rules:
                        s = tok.setdefault(str(cf.sqref), set())
                        for f in r.formula or []:
                            s.update(x for x in lits(f) if x)
                        if r.text:
                            s.add(r.text)
            for rg in sorted(set(tok_es) | set(tok_en)):
                e, m = tok_es.get(rg, set()), tok_en.get(rg, set())
                n['rangos'] += 1
                n['tokens'] += len(m)
                if len(e) != len(m) or len({x.lower() for x in m}) != len(m):
                    R.fallo('%s:%s CF %s: %d tokens ES → %d EN (colisión) %r' % (en, ws_en.title, rg, len(e), len(m), sorted(m)))
                fuera = sorted(x for x in m if mapas.necesita_clave(x) and x not in CLAVES_EN)
                if fuera:
                    R.fallo('%s:%s CF %s: tokens ∉ CLAVES %r' % (en, ws_en.title, rg, fuera))
                for a in m:
                    for b in m:
                        if a != b and a.lower() in b.lower():
                            esa = [k for k, v in CLAVES.items() if v == a]
                            esb = [k for k, v in CLAVES.items() if v == b]
                            if not any(x.lower() in y.lower() for x in esa for y in esb):
                                R.fallo('%s:%s CF %s: token %r contenido en %r (colisión SEARCH)' % (en, ws_en.title, rg, a, b))
    for e in mapas.cruzar_censo(os.path.join(AQUI, 'censo_es.json')):
        R.fallo('mapas.cruzar_censo: ' + e)
    for k, v in sorted(n.items()):
        R.dato(k, v)


# ==========================================================================
# G6 — cálculo
# ==========================================================================
def celdas_sin_v(path):
    fuera = set()
    with zipfile.ZipFile(path) as z:
        wbx = z.read('xl/workbook.xml').decode('utf-8')
        rels = z.read('xl/_rels/workbook.xml.rels').decode('utf-8')
        destino = {}
        for tag in re.findall(r'<Relationship\b[^>]*>', rels):
            i, t = re.search(r'\bId="([^"]+)"', tag), re.search(r'\bTarget="([^"]+)"', tag)
            if i and t:
                destino[i.group(1)] = t.group(1)
        for tag in re.findall(r'<sheet\b[^>]*>', wbx):
            nombre = html.unescape(re.search(r'\bname="([^"]*)"', tag).group(1))
            rid = re.search(r'\br:id="([^"]+)"', tag).group(1)
            t = destino[rid]
            parte = t.lstrip('/') if t.startswith('/') else 'xl/' + t
            xml = z.read(parte).decode('utf-8')
            for m in re.finditer(r'<c\b([^>]*?)(?:/>|>(.*?)</c>)', xml, re.S):
                dentro = m.group(2) or ''
                if '<f' in dentro and not re.search(r'<v>[^<]', dentro) and '<v></v>' not in dentro \
                        and '<v/>' not in dentro and '<v />' not in dentro:
                    fuera.add((nombre, re.search(r'\br="([A-Z]+\d+)"', m.group(1)).group(1)))
    return fuera


def recalcular(path, fn_mutar=None):
    """Copia de `path` → (opcional) mutación con openpyxl → inject_cache con IRR → TIR comprobada → wb data_only."""
    tmp = tempfile.mkdtemp(prefix='fin-calc-')
    try:
        p = os.path.join(tmp, os.path.basename(path))
        shutil.copy2(path, p)
        if fn_mutar is not None:
            wb = openpyxl.load_workbook(p)
            fn_mutar(wb)
            wb.save(p)
        nf, inj, fail = A.inyectar_cache(p)
        irr = None
        if os.path.basename(path) == A.FICHERO_07:
            irr = A.cachear_irr(p)
        return openpyxl.load_workbook(p, data_only=True), fail, irr
    finally:
        shutil.rmtree(tmp)


def referencia(pes, corto):
    """Caché de REFERENCIA: el ES publicado con VALORES_EN y el patrón D10 aplicados (pestañas ES), recalculado."""
    def mutar(wb):
        for (c, h, ref), v in mapas.VALORES_EN.items():
            if c != corto:
                continue
            ws = wb[h]
            for coord in A.rango(ref):
                if len(list(A.rango(ref))) > 1 and ws[coord].value is None:
                    continue
                ws[coord].value = v
        for ws in wb.worksheets:
            for cel in ws._cells.values():
                if cel.data_type == 'f' and mapas.patron(cel.value) in mapas.FORMULAS_EN:
                    cel.value = mapas.FORMULAS_EN[mapas.patron(cel.value)].replace('{r}', str(cel.row))
    wb, fail, _irr = recalcular(pes, mutar)          # el 07 ES evalúa IRR con el parche de inyectar_cache
    return wb, fail


# Cifras de SPEC §5 que deben verse en la caché EN: (libro, hoja EN, celda) → valor
ESPERADOS = OrderedDict([
    (('02', 'Scenarios', 'D10'), 31460), (('02', 'Scenarios', 'D16'), 0.173204),
    (('02', 'Break-Even', 'C8'), 23076.92), (('02', 'Break-Even', 'C10'), 41), (('02', 'Break-Even', 'C16'), 24307.69),
    (('02', 'Break-Even', 'C17'), 43),
    (('03', 'Monthly Cash Flow', 'B5'), 15000), (('03', 'Monthly Cash Flow', 'M29'), 15000),
    (('03', 'Cash Alerts', 'C17'), 15000), (('03', 'Cash Alerts', 'E17'), 5000),
    (('06', 'Ratios', 'C6'), 31460), (('06', 'Ratios', 'C23'), 0.172994), (('06', 'Ratios', 'C18'), 0.28),
    (('06', 'Ratios', 'C30'), 31460 / 860.0),                                     # D14: ventas / 860 sq ft
    (('06', 'Ratios', 'E18'), CLAVES['✅ Excelente']),                             # D15: labor 28 % < 30 % objetivo US
    (('B08', 'Simulator', 'C14'), 31460), (('B08', 'Simulator', 'C20'), 0.172924), (('B08', 'Simulator', 'C21'), 65282.4),
    (('B08', 'Comparison', 'D4'), 377520), (('B08', 'Comparison', 'D7'), 65282.4),
    (('07', 'Projections', 'B24'), '—'),                                         # sin datos: sin TIR, como el ES
    (('07', 'Ratios', 'C5'), CLAVES['Indica el préstamo']),
])

# Caso trazado de la TIR (SPEC §8 G6): EBITDA = flujo / (1 − 25 %) → flujo libre del proyecto = −150,000 / 30,000 /
# 45,000 / 60,000 / 70,000 / 0 → 11.9592 %; préstamo de ejemplo SBA 7(a) 100,000 al 10 % y 10 años → DSCR 1.84×
TRAZA_INVERSION = 150000
TRAZA_FLUJOS = [30000, 45000, 60000, 70000, 0]
TRAZA_PRESTAMO = 100000
TIR_SPEC = 0.119592


def _irr_biseccion(fl, lo=-0.99, hi=1.0):
    """TIR por bisección pura (independiente de la Newton del motor)."""
    f = lambda r: sum(x / (1 + r) ** i for i, x in enumerate(fl))      # noqa: E731
    a, b = f(lo), f(hi)
    if a * b > 0:
        return None
    for _ in range(300):
        m = (lo + hi) / 2
        fm = f(m)
        if a * fm <= 0:
            hi, b = m, fm
        else:
            lo, a = m, fm
    return (lo + hi) / 2


def mutar_traza(wb, en=True):
    h = (lambda x: mapas.HOJAS[x]) if en else (lambda x: x)
    wb[h('Resumen Ejecutivo')]['C13'].value = TRAZA_INVERSION
    for col, fl in zip('BCDEF', TRAZA_FLUJOS):
        wb[h('Proyecciones')][col + '4'].value = fl / (1 - 0.25)
    wb[h('Financiación')]['C4'].value = TRAZA_PRESTAMO


def caso_tir(R, carpeta, n):
    pen = fichero_en(carpeta, '07')
    wv, fail, irr = recalcular(pen, lambda wb: mutar_traza(wb, True))
    wves, fail_es, irr_es = recalcular(fichero_es('07'), lambda wb: mutar_traza(wb, False))
    P, Rt, F = wv['Projections'], wv['Ratios'], wv['Loan Schedule']
    flujos = [P[c + '22'].value for c in 'BCDEFG']
    t, t_es = P['B24'].value, wves['Proyecciones']['B24'].value
    bis = _irr_biseccion([float(x) for x in flujos])
    tasa = P['C27'].value
    van_py = sum(x / (1 + tasa) ** i for i, x in enumerate(flujos))
    acum, pb = 0.0, None
    for i, x in enumerate(flujos):
        prev, acum = acum, acum + x
        if i and acum >= 0 > prev and pb is None:
            pb = (i - 1) + (-prev / x)
    cuota = TRAZA_PRESTAMO * 0.10 / (1 - 1.10 ** -10)
    dscr = 30000 / cuota
    R.dato('tir_traza', t)
    R.dato('tir_traza_es', t_es)
    R.dato('van_traza', P['B25'].value)
    R.dato('payback_traza', P['B26'].value)
    R.dato('dscr_traza', Rt['C5'].value)
    n['casos_tir'] += 1
    chequeos = [
        ('fallos pycel 0 (EN y ES)', fail == 0 and fail_es == 0),
        ('flujos = −150,000 / 30,000 / 45,000 / 60,000 / 70,000 / 0',
         all(abs(a - b) < 1e-6 for a, b in zip(flujos, [-TRAZA_INVERSION] + TRAZA_FLUJOS))),
        ('TIR = 11.9592 % (SPEC)', num(t) and abs(t - TIR_SPEC) < 5e-7),
        ('TIR EN = TIR ES', num(t) and num(t_es) and abs(t - t_es) < 1e-12),
        ('TIR = bisección independiente', num(t) and bis is not None and abs(t - bis) < 1e-8),
        ('TIR = cachear_irr (Python sobre la caché)', irr is not None and irr['tir_python'] is not None
         and abs(irr['tir_python'] - t) < 1e-12),
        ('VAN 8 % = Python', num(P['B25'].value) and abs(P['B25'].value - van_py) < 0.01),
        ('payback = Python (3.21 años)', num(P['B26'].value) and pb is not None and abs(P['B26'].value - pb) < 1e-9),
        ('cuota SBA 10 % / 10 años', num(F['C8'].value) and abs(F['C8'].value - cuota) < 0.01),
        ('intereses año 1 = 10,000', num(P['B13'].value) and abs(P['B13'].value - 10000) < 0.01),
        ('DSCR = 30,000 / cuota (1.84×)', num(Rt['C5'].value) and abs(Rt['C5'].value - dscr) < 1e-6),
        ('DSCR ≥ 1.25 → ✅', Rt['E5'].value == '✅'),
        ('resumen ejecutivo: payback', wv['Executive Summary']['C19'].value == P['B26'].value),
    ]
    for nombre, ok in chequeos:
        if not ok:
            R.fallo('TIR/DSCR caso trazado: %s · TIR %r (ES %r, bisección %r) · VAN %r · payback %r · DSCR %r · flujos %r'
                    % (nombre, t, t_es, bis, P['B25'].value, P['B26'].value, Rt['C5'].value, flujos))


def caso_d10(R, carpeta, n):
    def mutar(wb):
        ws = wb['Kitchen Equipment']
        ws['B5'].value, ws['H5'].value = 7000, 7
        ws['B6'].value, ws['H6'].value = 7000, None
    wv, fail, _ = recalcular(fichero_en(carpeta, '04'), mutar)
    ws = wv['Kitchen Equipment']
    R.dato('d10_7000_7', ws['I5'].value)
    R.dato('d10_7000_vacia', ws['I6'].value)
    n['casos_d10'] += 1
    if fail or not (num(ws['I5'].value) and abs(ws['I5'].value - 1000) < 1e-9) or ws['I6'].value != 0 \
            or not (num(wv['Summary']['G6'].value) and abs(wv['Summary']['G6'].value - 1000) < 1e-9):
        R.fallo('D10: 7,000 / 7 años → %r (1,000) · vida vacía → %r (0) · Summary G6 %r · fallos pycel %s'
                % (ws['I5'].value, ws['I6'].value, wv['Summary']['G6'].value, fail))


def gateG6(R, carpeta, datos, casos=True):
    R.gate('G6 cálculo')
    n = Counter()
    movidas = []
    for es, en, corto, pes, pen in libros(carpeta):
        wf, wv = cargar(pen), cargar(pen, data_only=True)
        wes_v = cargar(pes, data_only=True)
        ref, fail_ref = referencia(pes, corto)
        if fail_ref:
            R.fallo('%s: la referencia ES tiene %d fallos de pycel' % (es, fail_ref))
        sin_v = celdas_sin_v(pen)
        for h_es in wes_v.sheetnames:
            h_en = mapas.HOJAS[h_es]
            ws, ws_v, wr, we = wf[h_en], wv[h_en], ref[h_es], wes_v[h_es]
            for c in ws._cells.values():
                if c.data_type != 'f':
                    continue
                n['formulas'] += 1
                if (ws.title, c.coordinate) in sin_v:
                    R.fallo('%s:%s!%s sin caché (<f> sin <v>)' % (en, ws.title, c.coordinate))
                    continue
                v, vr, ve = ws_v[c.coordinate].value, wr[c.coordinate].value, we[c.coordinate].value
                if isinstance(v, str) and v in ERRORES:
                    R.fallo('%s:%s!%s = %s' % (en, ws.title, c.coordinate, v))
                    continue
                n['con_valor' if v is not None else 'vacias_legitimas'] += 1
                if vr != ve and not (num(vr) and num(ve) and abs(vr - ve) < TOL):
                    movidas.append('%s %s!%s %r→%r' % (corto, h_en, c.coordinate, ve, vr))
                if (vr is None) != (v is None):
                    R.fallo('%s:%s!%s caché %r (referencia %r)' % (en, ws.title, c.coordinate, v, vr))
                elif isinstance(vr, str) and vr in CLAVES:
                    n['veredictos'] += 1
                    if v != CLAVES[vr]:
                        R.fallo('%s:%s!%s veredicto %r ≠ CLAVES[%r] = %r' % (en, ws.title, c.coordinate, v, vr, CLAVES[vr]))
                elif isinstance(v, str):
                    n['textos_calculados'] += 1
                    for motivo, ctx in restos_espanol(v):
                        R.fallo('%s:%s!%s texto calculado con %s · …%s…' % (en, ws.title, c.coordinate, motivo, ctx))
                elif num(vr):
                    n['numeros'] += 1
                    if not (num(v) and abs(v - vr) < TOL):
                        R.fallo('%s:%s!%s número %r ≠ referencia %r' % (en, ws.title, c.coordinate, v, vr))
    R.dato('movidas_por_spec', len(movidas))
    R.dato('movidas_muestra', movidas[:12])
    for (corto, hoja, coord), e in ESPERADOS.items():
        v = cargar(fichero_en(carpeta, corto), True)[hoja][coord].value
        n['esperados'] += 1
        ok = (num(v) and num(e) and abs(v - e) < (0.00051 if abs(e) < 1 else 0.006 if isinstance(e, int) else 0.01)) \
            or v == e
        if not ok:
            R.fallo('%s %s!%s = %r ≠ SPEC %r' % (corto, hoja, coord, v, e))
    # la TIR cacheada de la entrega (sin datos) = la que da Python sobre los flujos cacheados
    fl, cache = A.leer_flujos(fichero_en(carpeta, '07'))
    t = A.tir(fl)
    if (t is None and cache != '—') or (t is not None and not (num(cache) and abs(cache - t) < 1e-9)):
        R.fallo('07 TIR cacheada %r ≠ Python %r (flujos %r)' % (cache, t, fl))
    if casos:
        caso_tir(R, carpeta, n)
        caso_d10(R, carpeta, n)
    for k, v in sorted(n.items()):
        R.dato(k, v)


# ==========================================================================
# G7 — reglas (SPEC §3 = mapas.REGLAS_US)
# ==========================================================================
def gateG7(R, carpeta, datos):
    R.gate('G7 reglas')
    n = Counter()
    RU = {k: v[0] for k, v in mapas.REGLAS_US.items()}
    for (corto, h_es, ref), v in mapas.VALORES_EN.items():
        ws = cargar(fichero_en(carpeta, corto))[mapas.HOJAS[h_es]]
        coords = list(A.rango(ref))
        wes = cargar(fichero_es(corto))[h_es]
        for coord in coords:
            if len(coords) > 1 and wes[coord].value is None:
                continue
            got = ws[coord].value
            ok = got == v or (num(got) and num(v) and abs(got - v) < 1e-9)
            n['valores'] += 1
            if not ok:
                R.fallo('VALORES_EN %s %s!%s = %r ≠ %r' % (corto, ws.title, coord, got, v))
    V = mapas.VALORES_EN
    vidas = {h: V[('04', h, 'H5:H%d' % r)] for h, r in (('Obra', 14), ('Equipamiento Cocina', 16),
                                                         ('Mobiliario Sala', 14), ('Tecnología', 12))}
    w7 = cargar(fichero_en(carpeta, '07'))
    comprobaciones = [
        ('sales tax 8 % (ejemplo)', V[('03', 'Parámetros', 'C7')] == 0.08 and '8 % = ejemplo' in RU['sales_tax']),
        ('sales tax sin crédito → 0', V[('03', 'Parámetros', 'C8')] == 0 == V[('03', 'Parámetros', 'C9')]
         and 'sin crédito' in RU['sales_tax']),
        ('tarjeta 80 %', V[('03', 'Parámetros', 'C4')] == 0.80 and RU['tarjeta'].startswith('80 %')),
        ('vidas útiles 10/7/7/5', [vidas[h] for h in ('Obra', 'Equipamiento Cocina', 'Mobiliario Sala', 'Tecnología')]
         == [10, 7, 7, 5] and 'vidas 10/7/7/5' in RU['amortizacion_libro']),
        ('licencias: vida vacía (gasto)', V[('04', 'Licencias', 'H5:H12')] is None),
        ('impuesto recuperable US 0 en el 04', all(v == 0 for (c, h, r), v in V.items() if c == '04' and r.startswith('C'))),
        ('860 sq ft', V[('06', 'Ratios', 'C11')] == 860 and '860 sq ft' in RU['sqft']),
        ('labor 30/35 %', (V[('06', 'Benchmarks', 'F5')], V[('06', 'Benchmarks', 'G5')]) == (0.30, 0.35)
         and 'labor 30-35 %' in RU['benchmarks']),
        ('ocupación 6/10 %', (V[('06', 'Benchmarks', 'F8')], V[('06', 'Benchmarks', 'G8')]) == (0.06, 0.10)
         and 'ocupación 6-10 %' in RU['benchmarks']),
        ('SBA 7(a) 10 % (ejemplo)', V[('07', 'Financiación', 'C5')] == 0.10 and '10 % = ejemplo' in RU['sba_tipo']),
        ('SBA 7(a) 10 años', V[('07', 'Financiación', 'C6')] == 10 and 'hasta 10 años' in RU['sba_plazo']),
        ('aportación propia SBA 10 %', V[('07', 'Ratios', 'G9')] == 0.10 and '10 %' in RU['aportacion_sba']),
        ('DSCR objetivo 1.25×', w7['Ratios']['F5'].value == 1.25 and '1.25' in RU['dscr_banco']),
        ('DSCR mínimo SBA 1.15×', w7['Ratios']['G5'].value == 1.15 and '1.15' in RU['dscr_sba']),
        ('impuesto efectivo 25 %', w7['Projections']['C29'].value == 0.25 and '21 %' in RU['impuesto_sociedades']),
        ('interest-only vacío / 0', w7['Loan Schedule']['C7'].value in (None, 0)),
    ]
    for nombre, ok in comprobaciones:
        n['reglas'] += 1
        if not ok:
            R.fallo('regla %s: xlsx/mapas ≠ SPEC §3' % nombre)
    # D10: el patrón de dotación en las 48 celdas y la DV de vida útil 0-50
    w4 = cargar(fichero_en(carpeta, '04'))
    for (c, h, sq), partida in mapas.DV_PARTIDAS.items():
        ws = w4[mapas.HOJAS[h]]
        for coord in A.rango_lista(partida[1]):
            fi = ws['I' + coord[1:]].value
            n['d10_formulas'] += 1
            if fi != '=IFERROR($B%s/$H%s,0)' % (coord[1:], coord[1:]):
                R.fallo('04 %s!I%s = %r (D10: =IFERROR(B/H,0))' % (ws.title, coord[1:], fi))
        dvs = [d for d in ws.data_validations.dataValidation if str(d.sqref) == partida[1]]
        if len(dvs) != 1 or (dvs[0].type, dvs[0].operator, dvs[0].formula1, dvs[0].formula2) != ('decimal', 'between', '0', '50'):
            R.fallo('04 %s: DV de vida útil %s ausente o distinta' % (ws.title, partida[1]))
    if n['d10_formulas'] != 48:
        R.fallo('D10: %d celdas de dotación ≠ 48' % n['d10_formulas'])
    for k, v in sorted(n.items()):
        R.dato(k, v)


# ==========================================================================
# G8 — censo-entregables y gate-no-latinos
# ==========================================================================
def gateG8(R, carpeta, datos):
    R.gate('G8 censo + no-latinos')
    real = os.path.abspath(carpeta) == os.path.abspath(A.DESTINO_REAL)
    cmds = [[sys.executable, CENSO_SCRIPT, '--only', mapas.SLUG if real else carpeta, '--fail'] + ([] if real else ['--letter']),
            [sys.executable, NO_LATINOS, '--only', mapas.SLUG if real else carpeta]]
    for cmd in cmds:
        r = subprocess.run(cmd, stdout=subprocess.PIPE, stderr=subprocess.STDOUT, universal_newlines=True, cwd=REPO)
        lineas = [x for x in r.stdout.strip().splitlines() if x.strip()]
        nombre = os.path.basename(cmd[1])
        R.dato(nombre, lineas[-1] if lineas else '')
        if r.returncode:
            R.fallo('%s exit %d: %s' % (nombre, r.returncode, ' | '.join(lineas[-6:])))


# ==========================================================================
# Autotest: cada gate tiene que cazar un defecto inyectado en una COPIA
# ==========================================================================
def _mutar(carpeta, fname, fn):
    p = os.path.join(carpeta, fname)
    if getattr(fn, 'zip', False):
        fn(p)
        return
    wb = openpyxl.load_workbook(p)
    fn(wb)
    wb.save(p)


def en_zip(fn):
    fn.zip = True
    return fn


def _reemplazar_zip(p, parte_rx, viejo, nuevo):
    with zipfile.ZipFile(p) as z:
        infos = z.infolist()
        partes = {i.filename: z.read(i.filename) for i in infos}
    hecho = False
    for k in partes:
        if re.fullmatch(parte_rx, k):
            x = partes[k].decode('utf-8')
            if viejo in x:
                partes[k] = x.replace(viejo, nuevo, 1).encode('utf-8')
                hecho = True
    if not hecho:
        raise RuntimeError('mutación sin efecto: %r' % viejo)
    with zipfile.ZipFile(p + '.tmp', 'w', zipfile.ZIP_DEFLATED) as zo:
        for i in infos:
            zo.writestr(i, partes[i.filename])
    os.replace(p + '.tmp', p)


def _set(ws, coord, v):
    ws[coord].value = v


def _lock(wb):
    from openpyxl.styles import Protection
    wb['Inputs']['C6'].protection = Protection(locked=True)


def _dv_titulo_id(wb):
    for dv in wb['Build-Out'].data_validations.dataValidation:
        if str(dv.sqref).startswith('H5'):
            dv.errorTitle = 'kitpf-v2 · useful_life'


def _cf_colision(wb):
    ws = wb['Checklist']
    for cf in ws.conditional_formatting:
        for r in cf.rules:
            r.formula = [f.replace('"In progress"', '"Pending"') for f in r.formula]
            if r.text == 'In progress':
                r.text = 'Pending'


def _formula_capex(wb):
    wb['Kitchen Equipment']['I5'].value = '=$B5*$H5'


@en_zip
def _grafico_ref(p):
    _reemplazar_zip(p, r'xl/charts/chart\d+\.xml', "'Break-Even'!$B$21:$B$28", "'Break-Even'!$B$22:$B$28")


@en_zip
def _grafico_es(p):
    _reemplazar_zip(p, r'xl/charts/chart\d+\.xml', '<a:t>Amount</a:t>', '<a:t>Importe (€)</a:t>')


@en_zip
def _tir_cache(p):
    """La caché de la TIR (Projections!B24) pasa de «—» a 0.5 sin tocar la fórmula."""
    with zipfile.ZipFile(p) as z:
        infos = z.infolist()
        partes = {i.filename: z.read(i.filename) for i in infos}
    n = 0
    for k in partes:
        if re.fullmatch(r'xl/worksheets/sheet\d+\.xml', k):
            x = partes[k].decode('utf-8')
            if 'IRR(' in x:
                x, n = re.subn(r'(<c r="B24"[^>]*>(?:(?!</c>).)*?)<v>[^<]*</v>', r'\1<v>0.5</v>', x, count=1, flags=re.S)
                partes[k] = x.encode('utf-8')
    if not n:
        raise RuntimeError('mutación sin efecto: TIR')
    with zipfile.ZipFile(p + '.tmp', 'w', zipfile.ZIP_DEFLATED) as zo:
        for i in infos:
            zo.writestr(i, partes[i.filename])
    os.replace(p + '.tmp', p)


MUTACIONES = [
    ('G1', '04-restaurant-startup-costs-budget.xlsx', _formula_capex, 'fórmula EN'),
    ('G1', '02-restaurant-break-even-calculator.xlsx', _lock, 'desbloqueadas'),
    ('G1', '02-restaurant-break-even-calculator.xlsx', _grafico_ref, 'referencias'),
    ('G2', '01-restaurant-financial-projections-3-year.xlsx',
     lambda wb: _set(wb['Instructions'], 'B7', 'Completa los datos mensuales de cada año'), 'ES'),
    ('G2', '06-restaurant-kpi-ratios-dashboard.xlsx', lambda wb: _set(wb['Instructions'], 'B6', 'Sales: 31,460 €'), 'moneda'),
    ('G2', '04-restaurant-startup-costs-budget.xlsx', _dv_titulo_id, 'ID interno'),
    ('G2', '03-restaurant-cash-flow-forecast.xlsx', lambda wb: _set(wb['Instructions'], 'B9', 'See file 07 (Restaurant '
                                                                  'Loan Proposal (Lender & Investor Summary)).'), 'cita anidada'),
    ('G2', '05-restaurant-pl-template-budget-vs-actual.xlsx', _grafico_es, 'moneda'),
    ('G2', '06-restaurant-kpi-ratios-dashboard.xlsx', lambda wb: _set(wb['Ratios'], 'B11', 'Dining room m²'), 'm²'),
    ('G3', '07-restaurant-loan-proposal-lender-summary.xlsx',
     lambda wb: setattr(wb['Collateral'].page_setup, 'paperSize', 9), 'paperSize'),
    ('G3', 'BONUS-09-pre-opening-financial-checklist.xlsx',
     lambda wb: setattr(wb['Checklist']['E5'], 'number_format', 'dd/mm/yyyy'), 'fecha'),
    ('G3', '05-restaurant-pl-template-budget-vs-actual.xlsx',
     lambda wb: setattr(wb['Jan']['D24'], 'number_format', '0.0'), 'formato'),
    ('G3', '01-restaurant-financial-projections-3-year.xlsx',
     lambda wb: setattr(wb['Year 2'].column_dimensions['A'], 'width', 28), 'ANCHOS'),
    ('G4', 'BONUS-09-pre-opening-financial-checklist.xlsx', lambda wb: _set(wb['Checklist'], 'F5', 'Done'), 'fuera de su DV'),
    ('G4', '04-restaurant-startup-costs-budget.xlsx', lambda wb: _set(wb['Technology'], 'H5', 60), 'fuera de su DV'),
    ('G5', 'BONUS-09-pre-opening-financial-checklist.xlsx', _cf_colision, 'colisión'),
    ('G6', '02-restaurant-break-even-calculator.xlsx', lambda wb: _set(wb['Inputs'], 'C20', 25), 'SPEC'),
    ('G6', '03-restaurant-cash-flow-forecast.xlsx', lambda wb: _set(wb['Monthly Cash Flow'], 'B5', 20000), 'SPEC'),
    ('G6', '07-restaurant-loan-proposal-lender-summary.xlsx', _tir_cache, 'TIR'),
    ('G7', '03-restaurant-cash-flow-forecast.xlsx', lambda wb: _set(wb['Assumptions'], 'C7', 0.1), 'VALORES_EN'),
    ('G7', '04-restaurant-startup-costs-budget.xlsx', lambda wb: _set(wb['Technology'], 'H6', 3), 'VALORES_EN'),
]
if mapas.AJUSTE_TEXTO:
    _c, _h, _rg = mapas.AJUSTE_TEXTO[0]

    def _sin_ajuste(wb):
        from openpyxl.styles import Alignment
        ws = wb[mapas.HOJAS[_h]]
        for coord in A.rango(_rg):
            if ws[coord].value is not None:
                ws[coord].alignment = Alignment(wrap_text=False)
                return
    MUTACIONES.append(('G3', mapas.FICHEROS[[e for e in mapas.LIBROS_ES if A.extraer_textos.corto(e) == _c][0]],
                       _sin_ajuste, 'ajuste de texto'))


def autotest(carpeta, mes, datos):
    ok = True
    base = ejecutar(carpeta, mes, ['G1', 'G2', 'G3', 'G4', 'G5', 'G6', 'G7'], datos, silencioso=True, casos=False)
    sucios = [g for g, v in base.gates.items() if v['fallos']]
    if sucios:
        print('AUTOTEST: la carpeta base no está limpia (%s): arregla eso primero' % sucios)
        return False
    for gate, fname, fn, aguja in MUTACIONES:
        tmp = tempfile.mkdtemp(prefix='fin-autotest-')
        try:
            for f in mapas.FICHEROS.values():
                shutil.copy2(os.path.join(carpeta, f), tmp)
            _mutar(tmp, fname, fn)
            if gate == 'G6' and not getattr(fn, 'zip', False):
                A.inyectar_cache(os.path.join(tmp, fname))
            R = ejecutar(tmp, mes, [gate], datos, silencioso=True, casos=False)
            fallos = [f for g in R.gates.values() for f in g['fallos']]
            cazado = any(aguja.lower() in f.lower() for f in fallos)
            ok &= cazado
            print('  autotest %s %-50s %-16s %s' % (gate, fname, aguja, 'CAZADO' if cazado else 'NO CAZADO %r' % fallos[:2]))
        finally:
            shutil.rmtree(tmp)
    print('  autotest: %d defectos inyectados' % len(MUTACIONES))
    return ok


# ==========================================================================
# main
# ==========================================================================
GATES = OrderedDict([('G1', gateG1), ('G2', gateG2), ('G3', gateG3), ('G4', gateG4), ('G5', gateG5),
                     ('G6', gateG6), ('G7', gateG7), ('G8', gateG8)])


def cargar_datos():
    return {'censo': json.load(open(os.path.join(AQUI, 'censo_es.json'), encoding='utf-8'))}


def ejecutar(carpeta, mes, solo, datos, silencioso=False, casos=True):
    R = Resultado()
    for nombre, fn in GATES.items():
        if solo and nombre not in solo:
            continue
        A.termica()
        if nombre == 'G3':
            fn(R, carpeta, datos, mes)
        elif nombre == 'G6':
            fn(R, carpeta, datos, casos=casos)
        else:
            fn(R, carpeta, datos)
        g = R.gates[R.actual]
        if not silencioso:
            estado = 'OK' if not g['fallos'] else 'FALLA (%d)' % len(g['fallos'])
            resumen = ' · '.join('%s=%s' % (k, v) for k, v in list(g['datos'].items())[:18] if not isinstance(v, list))
            print('%-24s %s  %s' % (R.actual, estado, resumen[:400]))
            for f in g['fallos'][:int(os.environ.get('FIN_MAX_FALLOS', '15'))]:
                print('    - ' + f[:260])
    return R


def main():
    ap = argparse.ArgumentParser(description='Gates F2 del Restaurant Financial Plan Kit Pro (EN)')
    ap.add_argument('--dir', default=A.DESTINO_REAL, help='carpeta con los 10 xlsx EN (por defecto dl/<slug>)')
    ap.add_argument('--mes', default=None, help='mes de la línea de versión (por defecto, el actual)')
    ap.add_argument('--solo', default=None, help='G1,G2,… (por defecto, todos)')
    ap.add_argument('--json', default=None)
    ap.add_argument('--autotest', action='store_true')
    args = ap.parse_args()
    mes = args.mes or A.mes_por_defecto()
    datos = cargar_datos()
    carpeta = os.path.abspath(args.dir)
    if args.autotest:
        ok = autotest(carpeta, mes, datos)
        print('AUTOTEST: %s' % ('OK' if ok else 'FALLA'))
        sys.exit(0 if ok else 1)
    solo = [s.strip().upper() for s in args.solo.split(',')] if args.solo else None
    R = ejecutar(carpeta, mes, solo, datos)
    fallos = sum(len(g['fallos']) for g in R.gates.values())
    if args.json:
        with open(args.json, 'w', encoding='utf-8') as fh:
            json.dump(R.gates, fh, ensure_ascii=False, indent=1, default=str)
    print('\n%s: %d gates, %d fallos (%s)' % ('VERDE' if not fallos else 'ROJO', len(R.gates), fallos, carpeta))
    sys.exit(1 if fallos else 0)


if __name__ == '__main__':
    main()
