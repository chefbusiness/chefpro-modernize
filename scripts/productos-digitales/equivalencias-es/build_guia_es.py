#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""build_guia_es.py — «Guía de Equivalencias para tu país»: un docx por cada uno de los 59 productos ES.

    python3 scripts/productos-digitales/equivalencias-es/build_guia_es.py [--only a,b] [--out-root DIR] [--check]

Spec: SPEC-BUILD.md. Reutiliza helpers de ../equivalents-guide/build_guide.py (EN) y el registro de productos de
censo-terminos-es.py. Python 3.7.
"""
import argparse
import collections
import datetime as dt
import importlib.util
import json
import math
import os
import re
import sys
import time

import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.abspath(os.path.join(HERE, '..', '..', '..'))


def _load(name, path):
    spec = importlib.util.spec_from_file_location(name, path)
    m = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(m)
    return m


EN = _load('build_guide_en', os.path.join(HERE, '..', 'equivalents-guide', 'build_guide.py'))
CENSO = _load('censo_es', os.path.join(HERE, 'censo-terminos-es.py'))

TEMPLATE = EN.TEMPLATE
DEFAULT_OUT = os.path.join(ROOT, 'astro-site', 'public', 'dl')
ZONA = os.path.join(ROOT, 'astro-site', 'src', 'lib', 'zona-app.ts')

INK, GOLD, GRAY = EN.INK, EN.GOLD, EN.GRAY
BRIGHT = RGBColor(0xFF, 0xD7, 0x00)
CODES = ['MX', 'GT', 'HN', 'SV', 'NI', 'CR', 'PA', 'CU', 'DO', 'PR', 'CO', 'VE', 'EC', 'PE', 'BO', 'CL', 'AR',
         'UY', 'PY', 'US', 'BR']
CAMPOS = ['impuesto_consumo', 'propina_servicio', 'autoridad_sanitaria', 'manipulador', 'licencias_apertura',
          'alcohol', 'forma_juridica', 'laboral', 'unidades_mercado', 'pagos_delivery', 'apoyos_emprendimiento',
          'tipos_local', 'adaptar_productos']
LEGALES = set(['impuesto_consumo', 'propina_servicio', 'autoridad_sanitaria', 'manipulador', 'licencias_apertura',
               'alcohol', 'forma_juridica', 'laboral', 'apoyos_emprendimiento'])
KIT_EXCEL = set(['kit-escandallos', 'pack-appcc', 'kit-gestion-personal', 'kit-inventario', 'kit-plan-financiero'])
ORIENTATIVO = 'Orientativo: confírmalo con la autoridad local.'
JG = 'https://johnguerrero.es'

# nombres de país por código (Cuba solo en el glosario)
PAIS_CU = 'Cuba'


# ---------------------------------------------------------------- datos
def tipo_de(slug):
    if slug.startswith('kit-tareas') or slug == 'mega-pack-tareas':
        return 'kit_tareas'
    if slug in KIT_EXCEL:
        return 'kit_excel'
    if slug.startswith('plan-'):
        return 'plan_negocio'
    if slug.startswith('guia-') or slug.startswith('manual-'):
        return 'guia_manual'
    if slug == 'pro-prompts-ebook':
        return 'ebook'
    raise ValueError('tipo desconocido: ' + slug)


def carpeta_de(slug, paths):
    if slug in ('mega-pack-tareas', 'pro-prompts-ebook'):
        return slug
    c = collections.Counter()
    for p in paths:
        parts = p.split('/')  # ['', 'dl', carpeta, fichero]
        if len(parts) >= 4:
            c[parts[2]] += 1
    if not c:
        raise ValueError('sin carpeta: ' + slug)
    return c.most_common(1)[0][0]


def product_labels():
    txt = open(ZONA, encoding='utf-8').read()
    out = {}
    for m in re.finditer(r"productId:\s*'([^']+)'[^\n]*?productLabel:\s*'((?:[^'\\]|\\.)*)'", txt):
        out[m.group(1)] = m.group(2).replace("\\'", "'")
    return out


class Data(object):
    def __init__(self):
        self.textos = json.load(open(os.path.join(HERE, 'textos.json'), encoding='utf-8'))
        self.terminos = json.load(open(os.path.join(HERE, 'terminos.json'), encoding='utf-8'))
        self.matriz = json.load(open(os.path.join(HERE, 'datos', 'matriz.json'), encoding='utf-8'))['productos']
        g = []
        for n in ('A', 'B'):
            g += json.load(open(os.path.join(HERE, 'glosario', n + '.json'), encoding='utf-8'))
        self.glosario = dict((x['id'], x) for x in g)
        self.fichas = {}
        import glob
        for p in glob.glob(os.path.join(HERE, 'fichas', '*.json')):
            f = json.load(open(p, encoding='utf-8'))
            self.fichas[f['codigo']] = f
        self.paises = dict((c, f['pais']) for c, f in self.fichas.items())
        self.paises['CU'] = PAIS_CU
        self.paises['BR'] = self.fichas['BR']['pais'] if 'BR' in self.fichas else 'Brasil'
        self.orden_fichas = sorted(self.fichas.values(), key=lambda f: _collate(f['pais']))
        es, en = CENSO.read_registry()
        self.registro = es
        self.labels = product_labels()


def _collate(s):
    tr = dict(zip('áéíóúüñ', 'aeiouun'))
    return ''.join(tr.get(ch, ch) for ch in s.lower())


# ---------------------------------------------------------------- contenido de sección 1/2
def variant_cell(t):
    g = t
    codes_order = dict((c, i) for i, c in enumerate(CODES))
    vs = [v for v in g['variantes'] if not v.get('pt')]
    pt = [v for v in g['variantes'] if v.get('pt')]
    parts = []
    for v in vs:
        cs = sorted(set(v['paises']), key=lambda c: codes_order.get(c, 99))
        parts.append('**%s** (%s)' % (v['termino'], ', '.join(cs)))
    if g.get('igual'):
        ig = sorted(set(g['igual']), key=lambda c: codes_order.get(c, 99))
        parts.append('Igual que en España: ' + ', '.join(ig))
    for v in pt:
        parts.append('Brasil (pt): **%s**' % v['termino'])
    return ' · '.join(parts)


def term_rows(D, ids):
    rows = []
    for tid in ids:
        g = D.glosario[tid]
        rows.append([g['es'], variant_cell(g), g.get('nota') or ''])
    return rows


def sec1_ids(D, slug):
    ap = D.matriz[slug]['terminos']
    tm = dict((t['id'], t) for t in D.terminos)
    ids = [tid for tid, n in ap.items() if n >= 1 and tid in D.glosario]
    return ids, ap, tm


def grouped(D, ids, key, tm):
    """[(tema, [ids])] en el orden de textos.temas."""
    by = collections.OrderedDict((t, []) for t in D.textos['temas'])
    for tid in ids:
        by.setdefault(tm[tid]['tema'], []).append(tid)
    out = []
    for tema, lst in by.items():
        if lst:
            lst.sort(key=key)
            out.append((tema, lst))
    return out


# ---------------------------------------------------------------- docx
def bm_name(code):
    return 'ficha_' + code


def add_bookmark(p, name, idn):
    st = OxmlElement('w:bookmarkStart')
    st.set(qn('w:id'), str(idn))
    st.set(qn('w:name'), name)
    en = OxmlElement('w:bookmarkEnd')
    en.set(qn('w:id'), str(idn))
    p._p.insert(1 if p._p.pPr is not None else 0, st)
    p._p.append(en)


def anchor_link(p, text, anchor, size=None):
    h = OxmlElement('w:hyperlink')
    h.set(qn('w:anchor'), anchor)
    h.set(qn('w:history'), '1')
    r = OxmlElement('w:r')
    rPr = OxmlElement('w:rPr')
    col = OxmlElement('w:color')
    col.set(qn('w:val'), '8A6D2B')
    u = OxmlElement('w:u')
    u.set(qn('w:val'), 'single')
    rPr.append(col)
    rPr.append(u)
    if size:
        sz = OxmlElement('w:sz')
        sz.set(qn('w:val'), str(int(size * 2)))
        rPr.append(sz)
    r.append(rPr)
    t = OxmlElement('w:t')
    t.text = text
    t.set(qn('xml:space'), 'preserve')
    r.append(t)
    h.append(r)
    p._p.append(h)


def estimate_pages(d):
    """Estimación: líneas por celda según anchura y suma de alturas, 50 líneas/página de texto."""
    lines = 0.0
    for p in d.paragraphs:
        n = max(1, int(math.ceil(len(p.text) / 95.0)))
        lines += n * 1.2
    for t in d.tables:
        ncols = len(t.columns)
        cpl = {1: 95, 2: 60, 3: 40}.get(ncols, 28)
        for r in t.rows:
            m = 1
            for c in r.cells:
                m = max(m, int(math.ceil(len(c.text) / float(cpl * (1.0 if ncols < 3 else 0.9)))))
            lines += m * 1.1 + 0.6
    pages = lines / 50.0
    # cada ficha empieza en página nueva: sumar el hueco de la última página de cada ficha
    return max(1, int(round(pages + 10 * 0.5)))


def build_one(D, slug, out_root):
    T = D.textos
    label = D.labels[slug]
    tipo = tipo_de(slug)
    carpeta = carpeta_de(slug, D.registro[slug])
    out = os.path.join(out_root, carpeta, 'guia-equivalencias-%s.docx' % slug)

    doc = docx.Document(TEMPLATE)
    body = doc.element.body
    for child in list(body):
        if child.tag != qn('w:sectPr'):
            body.remove(child)
    pkg = doc.part.package
    for r_id, rel in list(pkg.rels.items()):
        if rel.reltype.endswith('/thumbnail'):
            del pkg.rels[r_id]

    sec = doc.sections[0]
    sec.different_first_page_header_footer = True
    sec.footer.is_linked_to_previous = False
    fp = sec.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    EN.run(fp, 'Guía de Equivalencias · %s · aichef.pro  ·  ' % label, size=8, color=GRAY)
    EN.page_field(fp)

    A = WD_ALIGN_PARAGRAPH
    para, run, rich = EN.para, EN.run, EN.rich

    # ---- Portada
    for _ in range(5):
        doc.add_paragraph()
    p = para(doc, align=A.CENTER)
    run(p, 'BONUS · AI Chef Pro', 'Georgia', 16, GOLD, True)
    p = para(doc, align=A.CENTER, before=10)
    run(p, 'Guía de Equivalencias', 'Georgia', 30, INK, True)
    p = para(doc, align=A.CENTER, after=4)
    run(p, 'para tu país', 'Georgia', 26, GOLD, False)
    p = para(doc, align=A.CENTER)
    run(p, EN.RULE, 'Georgia', 14, GOLD, False)
    p = para(doc, align=A.CENTER, before=24)
    run(p, '%s · España ↔ Hispanoamérica, EE. UU. y Brasil' % label, 'Calibri', 14, GRAY, False)
    p = para(doc, align=A.CENTER, before=6)
    run(p, 'Incluida con tu %s' % label, 'Calibri', 12, INK, True)
    for _ in range(5):
        doc.add_paragraph()
    p = para(doc, align=A.CENTER)
    run(p, 'Edición octubre 2026', 'Calibri', 11, GRAY, False)
    p = para(doc, align=A.CENTER)
    EN.hyperlink(p, 'John Guerrero', JG, size=12)
    run(p, ', consultor gastronómico', 'Calibri', 12, INK, False)
    p = para(doc, align=A.CENTER)
    run(p, 'aichef.pro', 'Calibri', 11, GOLD, False)
    EN.page_break(doc)

    # ---- Cómo usar + Aviso + leyenda
    doc.add_paragraph('Cómo usar esta guía', style='Heading 1').paragraph_format.page_break_before = False
    para(doc, T['como_usar']).paragraph_format.line_spacing = EN.BODY_LS
    EN.note(doc, '**Aviso.** ' + T['aviso'])
    doc.add_paragraph('Códigos de país', style='Heading 2')
    leg = [c for c in CODES]
    while len(leg) % 3:
        leg.append(None)
    rows = [[leg[i], leg[i + 1], leg[i + 2]] for i in range(0, len(leg), 3)]
    legend_table(doc, rows, D)
    para(doc, 'Cuba aparece solo en el vocabulario de la guía: no tiene ficha propia en la sección 4.').paragraph_format.line_spacing = EN.BODY_LS

    # ---- Sección 1
    ids, ap, tm = sec1_ids(D, slug)
    section_h1(doc, '1. Tu %s en tu país' % label, first=False, new_page=True)
    para(doc, T['nota_tipo'][tipo]).paragraph_format.line_spacing = EN.BODY_LS
    for tema, lst in grouped(D, ids, lambda t: (-ap[t], D.glosario[t]['es']), tm):
        doc.add_paragraph(T['temas'][tema], style='Heading 2').paragraph_format.keep_with_next = True
        EN.table(doc, ['En el producto (España)', 'Cómo se dice en tu país', 'Ojo con'],
                 term_rows(D, lst), [1.4, 3.5, 1.4])

    # ---- Sección 2
    extra = [t['id'] for t in D.terminos if t.get('comun') and t['id'] not in set(ids) and t['id'] in D.glosario]
    section_h1(doc, '2. Otros términos que te encontrarás')
    for tema, lst in grouped(D, extra, lambda t: D.glosario[t]['es'].lower(), tm):
        doc.add_paragraph(T['temas'][tema], style='Heading 2').paragraph_format.keep_with_next = True
        EN.table(doc, ['En España', 'Cómo se dice en tu país', 'Ojo con'], term_rows(D, lst), [1.4, 3.5, 1.4])

    # ---- Sección 3
    section_h1(doc, '3. Medidas y unidades')
    para(doc, T['medidas_intro']).paragraph_format.line_spacing = EN.BODY_LS
    cv = T['conversiones']
    EN.table(doc, cv['head'], cv['rows'], [1.9, 1.6, 2.8])
    para(doc, 'Para las unidades de mercado propias de cada país, mira el campo «%s» de su ficha en la sección 4.'
         % T['campos']['unidades_mercado']).paragraph_format.line_spacing = EN.BODY_LS

    # ---- Sección 4: fichas
    first = True
    for f in D.orden_fichas:
        if first:
            h1 = doc.add_paragraph('4. Tu país, en una página', style='Heading 1')
            h1.paragraph_format.page_break_before = True
            h1.paragraph_format.keep_with_next = True
            ip = para(doc, 'Ir a: ')
            ip.paragraph_format.line_spacing = EN.BODY_LS
            for k, ff in enumerate(D.orden_fichas):
                if k:
                    EN.run(ip, ' · ', size=9.5)
                anchor_link(ip, ff['pais'], bm_name(ff['codigo']), size=9.5)
            h2 = doc.add_paragraph(f['pais'], style='Heading 2')
            first = False
        else:
            h2 = doc.add_paragraph(f['pais'], style='Heading 2')
            h2.paragraph_format.page_break_before = True
        h2.paragraph_format.keep_with_next = True
        add_bookmark(h2, bm_name(f['codigo']), 100 + D.orden_fichas.index(f))
        mes = f.get('consultado', '2026-10')
        try:
            y, m = mes.split('-')
            mes_txt = ['ene', 'feb', 'mar', 'abr', 'may', 'jun', 'jul', 'ago', 'sep', 'oct', 'nov', 'dic'][int(m) - 1] + '-' + y
        except Exception:
            mes_txt = mes
        p = para(doc)
        run(p, 'Moneda: %s · Datos consultados en %s' % (f['moneda'], mes_txt), size=9.5, color=GRAY)
        rows = []
        for c in CAMPOS:
            v = f['campos'][c]
            txt = v['texto']
            if c in LEGALES and v.get('verificado') is False:
                txt += ' _' + ORIENTATIVO + '_'
            rows.append([T['campos'][c], txt])
        EN.table(doc, ['Tema', 'Qué debes saber'], rows, [1.5, 4.8])
        if f.get('trampas'):
            doc.add_paragraph('Falsos amigos', style='Heading 3' if 'Heading 3' in [s.name for s in doc.styles] else 'Heading 2')
            EN.bullets(doc, f['trampas'])
        seen = set()
        srcs = []
        for c in CAMPOS:
            for s in f['campos'][c].get('fuentes', []):
                if s['url'] not in seen:
                    seen.add(s['url'])
                    srcs.append(s)
        if srcs:
            p = para(doc, after=2)
            p.paragraph_format.line_spacing = 1.0
            run(p, 'Fuentes: ', size=7.5, bold=True)
            for k, sr in enumerate(srcs):
                if k:
                    run(p, ' · ', size=7.5, color=GRAY)
                EN.hyperlink(p, sr['titulo'], sr['url'], size=7.5)

    # ---- Sección 5
    h1 = doc.add_paragraph('5. ¿Lo adaptamos por ti?', style='Heading 1')
    h1.paragraph_format.page_break_before = True
    para(doc, T['cta']['texto']).paragraph_format.line_spacing = EN.BODY_LS
    para(doc, 'Todos los productos: [aichef.pro/productos-digitales](https://aichef.pro/productos-digitales) · '
         'Consultoría: [johnguerrero.es](https://johnguerrero.es)')

    cp = doc.core_properties
    cp.title = 'Guía de Equivalencias para tu país — %s' % label
    cp.subject = 'Vocabulario, impuestos, sanidad, normas laborales y unidades de 20 países para adaptar tu producto'
    cp.author = 'AI Chef Pro'
    cp.last_modified_by = 'AI Chef Pro'
    cp.keywords = 'equivalencias, hostelería, vocabulario, Hispanoamérica, %s' % label
    cp.category = 'Guía bonus'
    cp.comments = 'Bonus incluido con tu producto. Edición octubre 2026.'
    cp.language = 'es-419'
    now = dt.datetime(2026, 10, 13, 12, 0, 0)
    cp.created = now
    cp.modified = now
    cp.revision = 1

    os.makedirs(os.path.dirname(out), exist_ok=True)
    doc.save(out)
    return out, len(ids), len(extra)


def section_h1(doc, title, first=False, new_page=False):
    h1 = doc.add_paragraph(title, style='Heading 1')
    h1.paragraph_format.page_break_before = new_page
    h1.paragraph_format.space_before = Pt(0 if new_page else 22)
    h1.paragraph_format.keep_with_next = True
    p = doc.add_paragraph()
    p.paragraph_format.keep_with_next = True
    EN.run(p, EN.RULE_H1, size=8, color=GOLD)


def legend_table(doc, rows, D):
    from docx.enum.table import WD_TABLE_ALIGNMENT
    t = doc.add_table(rows=0, cols=3)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    EN.borders(t)
    EN.cell_margins(t)
    for r in rows:
        row = t.add_row()
        EN.no_split(row)
        for i, c in enumerate(r):
            if c is None:
                continue
            p = row.cells[i].paragraphs[0]
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            EN.run(p, c + ' · ', size=EN.TABLE_PT, bold=True, color=INK)
            nm = D.paises[c]
            if c in D.fichas:
                anchor_link(p, nm, bm_name(c), size=EN.TABLE_PT)
            else:
                EN.run(p, nm + ' (solo vocabulario)', size=EN.TABLE_PT, color=INK)
            if c == 'BR':
                EN.run(p, ' (portugués)', size=EN.TABLE_PT, color=INK)
    for row in t.rows:
        for i in range(3):
            row.cells[i].width = docx.shared.Inches(2.1)
    for i, gc in enumerate(t._tbl.tblGrid.findall(qn('w:gridCol'))):
        gc.set(qn('w:w'), str(int(2.1 * 1440)))
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


# ---------------------------------------------------------------- patch de EN.table: cabecera #111 + dorado
def _patched_table(doc, head, rows, widths, first_bold=True):
    from docx.enum.table import WD_TABLE_ALIGNMENT
    t = doc.add_table(rows=1, cols=len(head))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    EN.borders(t)
    EN.cell_margins(t)
    hdr = t.rows[0]
    EN.repeat_header(hdr)
    for i, h in enumerate(head):
        c = hdr.cells[i]
        EN.shade(c, '111111')
        EN.cell_text(c, h, bold_all=True, color=BRIGHT)
    for n, r in enumerate(rows):
        row = t.add_row()
        EN.no_split(row)
        for i, val in enumerate(r):
            c = row.cells[i]
            val = val or ''
            if i == 0 and first_bold:
                EN.shade(c, 'F3EBD8')
                EN.cell_text(c, val, bold_all=True, color=INK)
            else:
                if n % 2:
                    EN.shade(c, 'FAF7F0')
                _cell_rich(c, val)
    for row in t.rows:
        for i, w in enumerate(widths):
            row.cells[i].width = docx.shared.Inches(w)
    for i, gc in enumerate(t._tbl.tblGrid.findall(qn('w:gridCol'))):
        gc.set(qn('w:w'), str(int(widths[i] * 1440)))
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t


def _cell_rich(cell, text):
    """Como cell_text, pero soporta _cursiva_ final (aviso «Orientativo»)."""
    m = re.match(r'^(.*?)\s*_(' + re.escape(ORIENTATIVO) + r')_$', text, re.S)
    if not m:
        EN.cell_text(cell, text, color=INK)
        return
    EN.cell_text(cell, m.group(1), color=INK)
    p = cell.paragraphs[0]
    EN.run(p, ' ' + m.group(2), size=EN.TABLE_PT, color=GRAY, italic=True)


EN.table = _patched_table


# ---------------------------------------------------------------- check
NO_LAT = re.compile(r'[Ѐ-ӿ֐-׿؀-ۿ฀-๿぀-ヿ㐀-鿿가-힯　-〿]')
VOSOTROS = re.compile(r'\b(vosotros|vosotras|os|vuestro|vuestra|vuestros|vuestras)\b', re.I)
VOSEO = re.compile(r'\b(tenés|podés|querés|sabés|vení|mirá|fijate|hacé|usá|elegí|sos|vos)\b', re.I)
MUST_LINKS = ['https://johnguerrero.es', 'mailto:info@aichef.pro', 'https://aichef.pro/productos-digitales']


def all_text(d):
    parts = [p.text for p in d.paragraphs]
    for t in d.tables:
        for r in t.rows:
            for c in r.cells:
                parts.append(c.text)
    for s in d.sections:
        parts += [p.text for p in s.footer.paragraphs]
    xml = d.element.xml
    return '\n'.join(parts) + '\n' + ' '.join(re.findall(r'<w:t[^>]*>([^<]*)</w:t>', xml))


def fixed_texts(D):
    T = D.textos
    out = [T['como_usar'], T['aviso'], T['medidas_intro'], T['cta']['texto']] + list(T['nota_tipo'].values())
    return out


def check_one(D, slug, out_root):
    errs = []
    carpeta = carpeta_de(slug, D.registro[slug])
    path = os.path.join(out_root, carpeta, 'guia-equivalencias-%s.docx' % slug)
    if not os.path.exists(path):
        return ['no existe ' + path], None
    if os.path.getsize(path) <= 25 * 1024:
        errs.append('pesa %d bytes (<= 25 KB)' % os.path.getsize(path))
    d = docx.Document(path)
    txt = all_text(d)
    m = NO_LAT.search(txt)
    if m:
        errs.append('no latino: %r' % txt[max(0, m.start() - 30):m.end() + 30])
    if '](' in txt or '**' in txt:
        errs.append('marcado sin convertir')
    targets = [r.target_ref for r in d.part.rels.values() if r.reltype.endswith('/hyperlink')]
    for must in MUST_LINKS:
        if must not in targets:
            errs.append('falta el enlace ' + must)
    # sección 1: ≥5 términos y cada uno con apariciones ≥ 1
    ids, ap, tm = sec1_ids(D, slug)
    if len(ids) < 5:
        errs.append('sección 1 con %d términos (<5)' % len(ids))
    s1 = _section1_terms(d)
    esmap = dict((D.glosario[t]['es'], t) for t in D.glosario)
    for es in s1:
        tid = esmap.get(es)
        if tid is None or D.matriz[slug]['terminos'].get(tid, 0) < 1:
            errs.append('término de sección 1 sin apariciones: %r' % es)
    if len(s1) != len(ids):
        errs.append('sección 1 muestra %d términos, esperados %d' % (len(s1), len(ids)))
    for c, nm in D.paises.items():
        if c in ('CU',):
            continue
        if nm not in txt:
            errs.append('falta ficha de ' + nm)
    n_h2 = sum(1 for p in d.paragraphs if p.style.name == 'Heading 2' and p.text in [f['pais'] for f in D.fichas.values()])
    if n_h2 != len(D.fichas):
        errs.append('fichas H2: %d de %d' % (n_h2, len(D.fichas)))
    for s in fixed_texts(D):
        for rx, nm in ((VOSOTROS, 'vosotros/os'), (VOSEO, 'voseo')):
            mm = rx.search(s)
            if mm:
                errs.append('%s en texto fijo: %r' % (nm, s[max(0, mm.start() - 30):mm.end() + 30]))
    for p in d.paragraphs:
        if p.style.name in ('Heading 1', 'Heading 2'):
            continue
    return errs, {'terminos': len(ids), 'paginas': estimate_pages(d), 'bytes': os.path.getsize(path)}


def _section1_terms(d):
    """Texto ES (columna 1) de las tablas entre H1 «1.» y H1 «2.»."""
    body = d.element.body
    in1 = False
    out = []
    from docx.table import Table
    from docx.text.paragraph import Paragraph
    for el in body.iterchildren():
        if el.tag == qn('w:p'):
            p = Paragraph(el, d)
            if p.style.name == 'Heading 1':
                in1 = p.text.startswith('1. ')
        elif el.tag == qn('w:tbl') and in1:
            t = Table(el, d)
            for r in t.rows[1:]:
                out.append(r.cells[0].text)
    return out


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--only')
    ap.add_argument('--check', action='store_true')
    ap.add_argument('--out-root', default=DEFAULT_OUT)
    a = ap.parse_args()
    D = Data()
    slugs = list(D.registro.keys())
    if len(slugs) != 59:
        print('ERROR: %d productos ES (esperados 59)' % len(slugs))
        return 1
    if a.only:
        sel = [s.strip() for s in a.only.split(',') if s.strip()]
        bad = [s for s in sel if s not in D.registro]
        if bad:
            print('slugs desconocidos:', bad)
            return 1
        slugs = sel
    for s in slugs:
        if s not in D.labels:
            print('sin productLabel:', s)
            return 1
        if s not in D.matriz:
            print('sin matriz:', s)
            return 1
    if not a.check:
        for s in slugs:
            t0 = time.time()
            out, n1, n2 = build_one(D, s, a.out_root)
            print('OK %s · sec1=%d sec2=%d · %.1f s · %s' % (s, n1, n2, time.time() - t0, os.path.relpath(out, ROOT)))
    failed = 0
    for s in slugs:
        errs, info = check_one(D, s, a.out_root)
        if errs:
            failed += 1
            print('FAIL %s' % s)
            for e in errs:
                print('   -', e)
        else:
            print('PASS %s · términos sec.1: %d · páginas ~%d · %d KB' % (
                s, info['terminos'], info['paginas'], info['bytes'] // 1024))
    print('%d/%d OK' % (len(slugs) - failed, len(slugs)))
    return 1 if failed else 0


if __name__ == '__main__':
    sys.exit(main())
