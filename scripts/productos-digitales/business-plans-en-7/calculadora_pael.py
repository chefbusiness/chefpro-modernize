#!/usr/bin/env python3
"""
calculadora_pael.py — Paella Catering Pricing Calculator (Pans, Rice and Price per Guest to Quote), D123-D124 del Paella
Catering Business Plan Kit (EN). Esqueleto, cobro, cancelación, service charge, sales tax + interruptor, margen calculado y
autotest de `../business-plans-en-6/calculadora_bbq.py` (con lo que cazó su revisión final: % con decimal donde importa,
«mean» de BLS, rótulos de precio que no confunden, unidades US con la métrica al lado, aviso del coste provisional);
escalado por paellera, paelleros = paelleras ÷ 3 y ayudante desde 4 paelleras u 80 invitados de
`../paellero-2_2/calculadora_pae_es.py` y del manual ES (Kit estimate). Entradas verdes con validación, salidas
calculadas, ejemplos con las mismas fórmulas, docProps D20, US Letter, sin anclas de precio. Cinco pestañas:

  1. Pricing Calculator            paquete (drop-off paella · live paella party · full-service live paella), menu style (8,
                                   de la hoja 2), invitados y guest minimum, tamaño de paellera (lista de la hoja 2), oz de
                                   arroz crudo por invitado (3.25), interruptor «Tapas hour?» con su coste por invitado,
                                   ensalada y pan, desechables, propano; personal (paella chefs, ayudante, servers, horas,
                                   sueldos W-2); alquileres, millas; margen objetivo, service charge, sales tax + interruptor,
                                   depósito del 50 % (Kit estimate: los caterers US piden 25-50 %) y 4 tramos de
                                   cancelación → paelleras, arroz, caldo, propano y cilindros de 20 lb, coste, precio por
                                   invitado, total, depósito, saldo, retenciones y margen calculado. Mientras falte algún
                                   precio de ingrediente del estilo, el coste de comida usa el Kit estimate por invitado
                                   del estilo (research §5: $5 chicken & chorizo · $8.50 seafood · $16.50 lobster) y lo dice.
  2. Paella Yield & Pan Sizing     7 tamaños 18-40 in (cm al lado): invitados por paellera (Kit: 36 in = 40 · 32 in = 30 ·
                                   26 in = 18; el resto, tabla ES convertida), aros del quemador, arroz y caldo por paellera
                                   (lb y gal, con kg y litros), propano por cocción; 8 menu styles con las oz por invitado
                                   de cada ingrediente, el coste Kit por invitado, el coste con TUS precios ($/lb en blanco),
                                   el MÍNIMO del Food Code y el hot holding en columnas propias, alérgenos típicos con
                                   «verify with your recipe and suppliers».
  3. Shopping List by Guest Count  tu evento + 40 / 80 / 150 / 300 (editables): paelleras y quemadores, arroz, caldo,
                                   ingredientes del estilo, azafrán (g), limones, pan, ensalada, alioli, propano y
                                   cilindros, platos, hielo y coolers para el marisco, chafing dishes de las tapas, equipo.
  4. 3 Example Quotes              drop-off corporativo 30 inv. · backyard paella party 45 · boda 110 con tapas hour y
                                   personal: a ±10 % del precio de su línea en la mezcla de research §8 ($22 / $34 / $52),
                                   sin ningún precio de ingrediente tecleado.
  5. Instructions

Corre en el VPS (openpyxl + pycel). Sesión Claude Code, 24-oct-2026 (F2, paso 1).

    python3 bp7.py calculadora              # genera + inject_cache + autotest
    python3 bp7.py calculadora --autotest   # solo el autotest sobre el fichero ya generado
"""
import argparse
import math
import os
import re
import sys
from collections import OrderedDict

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill, Protection
from openpyxl.worksheet.datavalidation import DataValidation

AQUI = os.path.dirname(os.path.abspath(__file__))
PD = os.path.dirname(AQUI)
REPO = os.path.abspath(os.path.join(AQUI, '..', '..', '..'))
sys.path.insert(0, os.path.join(PD, 'business-plans-en'))
sys.path.insert(0, PD)
os.environ.setdefault('BP_CONJUNTO', 'pael')
import mapas as M                                                                  # noqa: E402

SLUG = 'paella-catering-business-plan'
FICHERO = 'paella-catering-pricing-calculator.xlsx'
TITULO = 'Paella Catering Pricing Calculator (Pans, Rice and Price per Guest to Quote)'
DESTINO = os.path.join(REPO, 'astro-site', 'public', 'dl', SLUG, FICHERO)
H1, H2, H3, H4, H5 = ('Pricing Calculator', 'Paella Yield & Pan Sizing', 'Shopping List by Guest Count',
                      '3 Example Quotes', 'Instructions')
HOJAS_CALCULO = (H1, H2, H3, H4)
VERDE = PatternFill('solid', fgColor='E8F5E9')
GRIS = PatternFill('solid', fgColor='F2F2F2')
F = 'Calibri'
MONEDA, ENTERO, PCT, PCT1, DEC1, DEC2 = '#,##0.00', '#,##0', '0%', '0.0%', '#,##0.0', '#,##0.00'
LB_KG, GAL_L = 0.4536, 3.785

DROP = 'Drop-off paella (cooked at the commissary, delivered hot)'
PARTY = 'Live paella party (cooked on site, salad & bread)'
FULL = 'Full-service live paella (paella chef & staff; weddings & corporate)'
# paquete → (invitados por server o buffet attendant (0 = ninguno), nota). Servers: «roughly one team member per 25
# guests» (miamigrillcatering.com, leída) en full-service; el resto, Kit estimate
PAQUETES = OrderedDict([
    (DROP, (0, 'Cooked at the commissary in the paella pans and delivered hot in insulated carriers, no servers on site. '
               'Keep it at 135 °F (57 °C) or above, or label the time it left hot holding and discard it 4 hours later '
               '(time as a public health control, only if your county accepts it).')),
    (PARTY, (0, 'Cooked live on site by the paella chef, with salad and bread. Guests serve themselves or the paella chef '
                'serves from the pan: no servers by default.')),
    (FULL, (25, 'Paella chef, helper and servers for weddings and corporate events: about one team member per 25 guests '
                '(miamigrillcatering.com).')),
])
NOMBRES_PAQ = list(PAQUETES)

# --------------------------------------------------------------------------
# 2. Paella Yield & Pan Sizing
#    A) 7 tamaños: (rótulo, cm, invitados por paellera [Kit], aros del quemador, propano por cocción lb [Kit], nota)
#       36 in = 40 · 32 in = 30 · 26 in = 18 (research §5, Kit); el resto, la tabla ES (manual técnico 1.2, tabla 1) × 1.33
#       (la misma relación que en 26-36 in). Propano: manual ES 1-1.5 kg por paellera de 90 cm (≈ 3 lb), escalado por área
# --------------------------------------------------------------------------
TAMANOS = [
    ('18 in (46 cm)', 46, 8, 1, 0.8, 'Side dish or a small backyard party.'),
    ('22 in (55 cm)', 55, 12, 2, 1.2, ''),
    ('26 in (65 cm)', 65, 18, 2, 1.6, 'Research figure (Kit estimate): about 18 guests.'),
    ('28 in (70 cm)', 70, 21, 2, 1.8, ''),
    ('32 in (80 cm)', 80, 30, 3, 2.4, 'Research figure (Kit estimate): about 30 guests.'),
    ('36 in (90 cm)', 90, 40, 3, 3.0, 'Research figure (Kit estimate): about 40 guests; a retailer says 40 to 50.'),
    ('40 in (100 cm)', 100, 50, 3, 3.7, 'Needs an XL 3-ring burner and wind guards; check that your burner fits it.'),
]
NOMBRES_TAM = [t[0] for t in TAMANOS]
Y_STOCK = 5                                                   # fila del parámetro «qt de caldo por lb de arroz»
Y_A0 = 9                                                      # primera fila de tamaños
Y_ACOL = OrderedDict([('cm', 'B'), ('guests', 'C'), ('rings', 'D'), ('rice', 'E'), ('rice_kg', 'F'), ('stock', 'G'),
                      ('stock_l', 'H'), ('propane', 'I'), ('propane_kg', 'J'), ('nota', 'K')])
T_TAM = "'%s'!$A$%d:$A$%d" % (H2, Y_A0, Y_A0 + len(TAMANOS) - 1)


def _tcol(col):
    return "'%s'!$%s$%d:$%s$%d" % (H2, col, Y_A0, col, Y_A0 + len(TAMANOS) - 1)


#    B) 11 ingredientes (oz por invitado; azafrán en g) × 8 menu styles, Kit estimate. Columnas B-L de la tabla de estilos
INGREDIENTES = [
    ('Chicken thighs, boneless (oz)', 'lb'), ('Spanish chorizo, cured (oz)', 'lb'), ('Rabbit (oz)', 'lb'),
    ('Shrimp (oz)', 'lb'), ('Mussels & clams, in shell (oz)', 'lb'), ('Squid or cuttlefish (oz)', 'lb'),
    ('Lobster, in shell (oz)', 'lb'), ('Firm white fish (oz)', 'lb'),
    ('Vegetables: green beans, lima beans, peppers, artichokes (oz)', 'lb'),
    ('Sofrito base: tomato, onion, garlic, olive oil (oz)', 'lb'), ('Saffron (g)', 'g'),
]
ING_COLS = 'BCDEFGHIJKL'
SEAFOOD_IDX = (3, 4, 5, 6, 7)                                 # shrimp, mussels & clams, squid, lobster, fish
_FC_MAR = '145 °F (63 °C) for 15 s: fish and shellfish'
_FC_AVE = '165 °F (74 °C): chicken'
_HH = '135 °F (57 °C) or above, or time as a public health control: 4 hours, then discard'
_VER = ' Verify with your recipe and suppliers.'
_MOL = 'mollusks (not a major US allergen, declare them anyway)'
ESTILOS = [
    ('Seafood (paella marinera)', [0, 0, 0, 2.0, 3.0, 1.5, 0, 1.0, 0.5, 1.5, 0.08], 8.50, _FC_MAR,
     'Crustacean shellfish, fish, %s; sulfites in some frozen shrimp.%s' % (_MOL, _VER)),
    ('Chicken & seafood', [2.5, 0, 0, 1.5, 2.0, 0.75, 0, 0, 1.0, 1.5, 0.08], 7.00, _FC_AVE + '; ' + _FC_MAR,
     'Crustacean shellfish, %s; sulfites in some frozen shrimp.%s' % (_MOL, _VER)),
    ('Chicken & chorizo', [3.5, 1.0, 0, 0, 0, 0, 0, 0, 1.5, 1.5, 0.08], 5.00,
     _FC_AVE + '. Spanish cured chorizo is ready to eat; raw Mexican-style chorizo is ground meat: 155 °F (68 °C) for 17 s',
     'Check the chorizo label (some brands add milk, soy or wheat).%s' % _VER),
    ('Valencian (chicken; rabbit on request)', [3.0, 0, 1.5, 0, 0, 0, 0, 0, 2.5, 1.0, 0.08], 6.00,
     _FC_AVE + '. Rabbit: check your county rules (cooking it with the chicken to 165 °F (74 °C) covers it)',
     'None of the 9 major allergens in a classic recipe.%s' % _VER),
    ('Farm vegetable (vegan)', [0, 0, 0, 0, 0, 0, 0, 0, 5.0, 2.0, 0.08], 4.00,
     '135 °F (57 °C): plant foods cooked for hot holding',
     'Check the vegetable stock base (some add celery, soy or wheat).%s' % _VER),
    ('Paella negra (squid ink)', [0, 0, 0, 1.0, 0, 3.0, 0, 0, 0.5, 1.5, 0], 6.50, _FC_MAR,
     'Crustacean shellfish, %s (squid and its ink).%s' % (_MOL, _VER)),
    ('Lobster & seafood (premium)', [0, 0, 0, 1.5, 2.0, 1.0, 4.0, 0, 0.5, 1.5, 0.08], 16.50, _FC_MAR,
     'Crustacean shellfish, %s.%s' % (_MOL, _VER)),
    ('Easy-eat peeled seafood', [0, 0, 0, 2.5, 0, 1.5, 0, 1.5, 0.5, 1.5, 0.08], 9.00, _FC_MAR,
     'Crustacean shellfish, fish, %s.%s' % (_MOL, _VER)),
]
NOMBRES_EST = [e[0] for e in ESTILOS]
Y_B_CAB = Y_A0 + len(TAMANOS) + 3                             # cabecera de la tabla de estilos (fila 19)
Y_B0 = Y_B_CAB + 1                                            # primera fila de estilos (20)
Y_PRECIO = Y_B0 + len(ESTILOS) + 1                            # fila de TUS precios por ingrediente (29)
Y_HELP_P = Y_PRECIO + 1                                       # precio por oz (o por g)
Y_HELP_F = Y_PRECIO + 2                                       # falta el precio (1 = sí)
Y_ARROZ = Y_PRECIO + 4                                        # precio del arroz bomba por lb (B33)
Y_BCOL = OrderedDict([('kit', 'M'), ('tuyo', 'N'), ('faltan', 'O'), ('fc', 'P'), ('hh', 'Q'), ('alerg', 'R')])
T_EST = "'%s'!$A$%d:$A$%d" % (H2, Y_B0, Y_B0 + len(ESTILOS) - 1)


def _ecol(col):
    return "'%s'!$%s$%d:$%s$%d" % (H2, col, Y_B0, col, Y_B0 + len(ESTILOS) - 1)


# --------------------------------------------------------------------------
# 1. Pricing Calculator: entradas (clave, rótulo, formato, valor por defecto, nota). Fila = 5 + índice
# --------------------------------------------------------------------------
ENTRADAS = [
    ('package', 'Package', '@', PARTY,
     'Choose from the list (table at the bottom of this tab): it sets the suggested servers. All three are priced per '
     'guest with the same formulas.'),
    ('style', 'Menu style', '@', 'Chicken & seafood',
     'Choose one of the 8 styles on the "%s" tab: it sets the ingredients per guest and the food cost.' % H2),
    ('guests', 'Guests expected', ENTERO, 58, 'Whole number. You cook and buy for these guests.'),
    ('minimum', 'Guest minimum in your contract', ENTERO, 20,
     'If fewer guests confirm, you bill the minimum (billable guests in the results). Operators read for this kit ask '
     'for 20 to 30 guests (tasteofpaella.com, zola.com). 0 = no minimum.'),
    ('pan_size', 'Paella pan size', '@', '36 in (90 cm)',
     'Choose from the list on the "%s" tab: it sets the guests per pan, the rice and stock per pan and the propane.' % H2),
    ('rice_oz', 'Uncooked rice per guest (oz)', DEC2, 3.25,
     'Kit estimate: 3 to 3.5 oz (85-100 g) of uncooked bomba rice per guest, a little more for big groups. 3.25 oz = '
     '92 g.'),
    ('tapas', 'Tapas hour?', '@', 'No',
     'Yes = a tapas hour before the paella (patatas bravas, Spanish tortilla, cheese and cold cuts, Serrano ham), costed '
     'per guest in the next row and priced with your margin.'),
    ('tapas_cost', 'Tapas food cost per guest', MONEDA, 4.50, 'Kit estimate: $3-6 per guest. Used only when the switch '
                                                              'above is Yes.'),
    ('salad_bread', 'Salad, bread, lemons and garlic aioli per guest', MONEDA, 1.25, 'Kit estimate: $1-1.50 per guest.'),
    ('disposables', 'Disposables per guest (plates, cutlery, napkins, delivery trays)', MONEDA, 0.75,
     'Kit estimate. Put 0 if the client rents dishes and flatware (then add the rental below).'),
    ('propane_price', 'Propane cost per lb', MONEDA, 1.75,
     'Kit estimate for 20-lb cylinder refills or exchanges ($10-15 for the 6-8 lb a 60-guest event burns). Use your own '
     'price.'),
    ('chefs', 'Paella chefs on event day', ENTERO, 1, 'Compare with the suggested number in the results.'),
    ('helpers', 'Helpers (prep cook or second cook)', ENTERO, 0, 'Compare with the suggested number in the results.'),
    ('servers', 'Servers and buffet attendants', ENTERO, 0, 'Compare with the suggested number in the results.'),
    ('event_hours', 'Hours each on event day (loading, travel, setup, cooking, service, breakdown)', DEC1, 6.0,
     'Kit estimate: arrive 2-3 hours before service; about 1 hour of cooking, 2 hours of service and 1 hour of '
     'breakdown (miamigrillcatering.com).'),
    ('prep_hours', 'Commissary prep hours for this event (sofrito, stock, seafood)', DEC1, 3.0,
     'Total hours the day before, paid at the helper wage. Kit estimate.'),
    ('wage_chef', 'Paella chef hourly wage', MONEDA, 22.0,
     'Kit estimate, about the 75th percentile of restaurant cooks (BLS OEWS, May 2025: mean $18.62, 75th percentile '
     '$21.45). W-2 employees; if you cook yourself, count your hours too.'),
    ('wage_helper', 'Helper hourly wage', MONEDA, 19.0,
     'Kit estimate, about the mean wage of restaurant cooks (BLS OEWS, May 2025: $18.62). W-2.'),
    ('wage_server', 'Server / attendant hourly wage', MONEDA, 18.0,
     'Kit estimate, about the 75th percentile of nonrestaurant food servers (BLS OEWS, May 2025: $18.51). W-2.'),
    ('payroll', "Employer payroll taxes and workers' comp (% of wages)", PCT, 0.10, 'Kit estimate.'),
    ('rentals', 'Rentals for this event (your cost; they carry your margin)', MONEDA, 0.0,
     'What the rental company charges you for tables, linens, dishes and flatware or a tent. They sit in the direct cost, '
     'so the target margin applies to them too.'),
    ('miles', 'Round-trip miles', ENTERO, 30, 'Commissary → venue → commissary.'),
    ('per_mile', 'Cost per mile (van)', MONEDA, 0.70,
     'Kit estimate: use your real cost per mile or the current IRS standard mileage rate.'),
    ('event_fee', 'Flat event or travel fee charged to the client', MONEDA, 0.0,
     'Optional flat charge on top of the per-guest price (some operators include travel up to 20 miles: '
     'cafesevilla.com). 0 = none.'),
    ('commission', 'Marketplace fee on this booking (% of subtotal)', PCT, 0.0,
     'Corporate catering marketplaces such as ezCater keep about 15% (Kit estimate). 0 = direct booking.'),
    ('margin', 'Target gross margin on the food & service subtotal', PCT, 0.50,
     'Share of the subtotal left after the direct costs above and the marketplace fee. Kit estimate.'),
    ('sc', 'Service charge (% of subtotal)', PCT, 0.0,
     'Added on top. The part you pay out to staff is wages, not tips (IRS Rev. Rul. 2012-18), with payroll taxes; the '
     'part you keep is revenue.'),
    ('tax', 'Sales tax rate (state + local)', PCT1, 0.08, 'Kit estimate: use your own state + local rate.'),
    ('tax_sc', 'Does your state tax the service charge?', '@', 'Yes',
     'Sales tax always applies to the food & service subtotal here. CA: a mandatory service charge is taxable. NY: not '
     'taxable only if it is separately stated, identified as a gratuity and all of it goes to the staff. TX: not taxable '
     'up to 20% if separately stated and paid to the staff. Check your state.'),
    ('deposit', 'Deposit to book the date (% of total)', PCT, 0.50,
     'Kit estimate: US caterers ask 25-50% to hold the date; the balance is due before the event. Put the same figure '
     'in your contract.'),
    ('d1', 'Cancellation tier 1 starts at (days before the event)', ENTERO, 60,
     'Tiers run with no gaps: tier 1 = this many days or more; tier 2 = from the next threshold up to here; and so on.'),
    ('d2', 'Cancellation tier 2 starts at (days before the event)', ENTERO, 30, ''),
    ('d3', 'Cancellation tier 3 starts at (days before the event)', ENTERO, 14, 'Tier 4 = fewer days than this, or a no-show.'),
    ('r1', 'Share of the total you keep in tier 1', PCT, 0.50, 'Kit estimate: the deposit.'),
    ('r2', 'Share of the total you keep in tier 2', PCT, 0.75, 'Kit estimate.'),
    ('r3', 'Share of the total you keep in tier 3', PCT, 0.90, 'Kit estimate.'),
    ('r4', 'Share of the total you keep in tier 4', PCT, 1.00, 'Kit estimate: the seafood is already bought.'),
]
K = {k: 5 + i for i, (k, *_r) in enumerate(ENTRADAS)}                  # clave → fila de la hoja 1
F_CALC = 5 + len(ENTRADAS) + 1                                         # cabecera del bloque de cálculo
SALIDAS = [
    ('billable', 'Billable guests', ENTERO, 'max({guests},{minimum})', 'The larger of guests expected and the minimum.'),
    ('guests_pan', 'Guests per pan for this size', ENTERO, 'tam:C', 'From the pan table on the "%s" tab.' % H2),
    ('pans', 'Paella pans to cook', ENTERO, 'roundup({guests}/{guests_pan})', 'Guests ÷ guests per pan, rounded up. '
                                                                             'One burner per pan.'),
    ('rice_lb', 'Uncooked rice (lb)', DEC1, '{guests}*{rice_oz}/16', 'Guests × oz per guest ÷ 16 (1 lb = 0.454 kg).'),
    ('stock_gal', 'Stock (gallons)', DEC1, 'caldo', 'Uncooked rice × the stock ratio on the "%s" tab (1 gallon = 3.8 '
                                                   'liters).' % H2),
    ('propane_lb', 'Propane for the cook (lb)', DEC1, 'propano', 'Pans × propane per cook for this size.'),
    ('cylinders', '20-lb propane cylinders to bring (plus one spare)', ENTERO, 'roundup({propane_lb}/20)+1',
     'Never run a pan on a nearly empty cylinder.'),
    ('missing', 'Ingredients of this style with no price', ENTERO, 'est:O', 'From the "%s" tab.' % H2),
    ('food_status', 'Ingredient prices', '@', 'estado',
     'While any price is missing, the food cost uses the style estimate. Type your prices on the "%s" tab.' % H2),
    ('style_cost', 'Paella food cost per guest for this style', MONEDA, 'coste_estilo',
     'With every price typed: your rice and ingredients per guest. Otherwise: the style estimate (Kit).'),
    ('c_food', 'Food cost (paella, salad, bread, lemons and aioli)', MONEDA, '{guests}*({style_cost}+{salad_bread})', ''),
    ('c_tapas', 'Tapas hour', MONEDA, 'if({tapas}=="Yes",{guests}*{tapas_cost},0)', ''),
    ('c_disp', 'Disposables', MONEDA, '{guests}*{disposables}', ''),
    ('c_propane', 'Propane', MONEDA, '{propane_lb}*{propane_price}', ''),
    ('c_staff', 'Staff cost (wages + payroll taxes)', MONEDA,
     '(({chefs}*{wage_chef}+{helpers}*{wage_helper}+{servers}*{wage_server})*{event_hours}+{prep_hours}*{wage_helper})*'
     '(1+{payroll})', ''),
    ('c_other', 'Rentals and mileage', MONEDA, '{rentals}+{miles}*{per_mile}', ''),
    ('c_total', 'Total direct cost', MONEDA, '{c_food}+{c_tapas}+{c_disp}+{c_propane}+{c_staff}+{c_other}', ''),
    ('c_guest', 'Direct cost per billable guest', MONEDA, '{c_total}/{billable}', ''),
    ('price', 'PRICE PER GUEST before fees, service charge and tax', MONEDA,
     'roundup({c_total}/(1-{margin}-{commission})/{billable}*2)/2',
     'Rounded up to the next 0.50. While "Ingredient prices" says Missing, this price uses the style food cost estimate, '
     'not your prices: treat it as provisional. Market check: live paella with salad and bread runs about $12-31 per '
     'person (zola.com, tasteofpaella.com) and a full Spanish menu with staff $44-65 (zola.com, miamigrillcatering.com).'),
    ('subtotal', 'Food & service subtotal', MONEDA, '{price}*{billable}+{event_fee}',
     'Price per guest × billable guests + event fee.'),
    ('sc_amt', 'Service charge', MONEDA, '{subtotal}*{sc}', ''),
    ('tax_amt', 'Sales tax', MONEDA, '({subtotal}+if({tax_sc}=="Yes",{sc_amt},0))*{tax}',
     'On the subtotal, plus the service charge if your state taxes it.'),
    ('total', 'TOTAL TO QUOTE', MONEDA, '{subtotal}+{sc_amt}+{tax_amt}', ''),
    ('dep_amt', 'Deposit to book the date', MONEDA, '{total}*{deposit}', ''),
    ('balance', 'Balance due before the event', MONEDA, '{total}-{dep_amt}', ''),
    ('comm_amt', 'Marketplace fee', MONEDA, '{subtotal}*{commission}', 'What the marketplace keeps.'),
    ('profit', 'Gross profit on the subtotal', MONEDA, '{subtotal}-{comm_amt}-{c_total}', ''),
    ('margin_calc', 'Gross margin (calculated)', PCT1, 'if({subtotal}==0,0,{profit}/{subtotal})',
     'Calculated from the costs above, never typed. It is at or above the target because the price is rounded up.'),
    ('chefs_sug', 'Suggested paella chefs (Kit estimate: one per 3 pans)', ENTERO, 'roundup({pans}/3)', ''),
    ('helper_sug', 'Suggested helper (Kit estimate: from 4 pans or 80 guests)', ENTERO,
     'if(or({pans}>=4,{guests}>=80),1,0)', ''),
    ('srv_sug', 'Suggested servers and attendants for this package', ENTERO, 'paq:2', 'From the package table below.'),
    ('tiers_check', 'Cancellation tiers check', '@', 'tramos', 'Tier days must go down: tier 1 > tier 2 > tier 3.'),
    ('ret1', 'Cancellation in tier 1: amount you keep', MONEDA, '{total}*{r1}', ''),
    ('ret2', 'Cancellation in tier 2: amount you keep', MONEDA, '{total}*{r2}', ''),
    ('ret3', 'Cancellation in tier 3: amount you keep', MONEDA, '{total}*{r3}', ''),
    ('ret4', 'Cancellation in tier 4 or no-show: amount you keep', MONEDA, '{total}*{r4}', ''),
]
S = {k: F_CALC + 1 + i for i, (k, *_r) in enumerate(SALIDAS)}
F_TABLA = F_CALC + 1 + len(SALIDAS) + 1                                # cabecera de la tabla de paquetes
_TP = lambda col: "'%s'!$%s$%d:$%s$%d" % (H1, col, F_TABLA + 1, col, F_TABLA + len(PAQUETES))   # noqa: E731
T_NOMBRES = _TP('A')

# 3 cotizaciones (Kit estimates; D124): cada una a ±10 % del precio de su línea en la mezcla de research §8, SIN precios
# de ingredientes (el coste de comida usa el Kit estimate por invitado del estilo)
OBJETIVO_EJ = OrderedDict([('Corporate drop-off, 30 guests', 22.0), ('Backyard paella party, 45 guests', 34.0),
                           ('Wedding, 110 guests, tapas hour + full staff', 52.0)])
_DEF = OrderedDict((k, v) for k, _r, _f2, v, _n in ENTRADAS)


def _ej(**kw):
    e = OrderedDict(_DEF)
    e.update(kw)
    return e


EJEMPLOS = OrderedDict([
    ('Corporate drop-off, 30 guests', _ej(
        package=DROP, style='Chicken & seafood', guests=30, minimum=20, chefs=1, helpers=0, servers=0, event_hours=3.0,
        prep_hours=2.0, rentals=0.0, miles=20, margin=0.40)),
    ('Backyard paella party, 45 guests', _ej(
        package=PARTY, style='Seafood (paella marinera)', guests=45, minimum=20, chefs=1, helpers=0, servers=0,
        event_hours=6.0, prep_hours=3.0, rentals=0.0, miles=30, margin=0.50)),
    ('Wedding, 110 guests, tapas hour + full staff', _ej(
        package=FULL, style='Chicken & seafood', guests=110, minimum=50, tapas='Yes', chefs=2, helpers=1, servers=5,
        event_hours=7.0, prep_hours=6.0, disposables=0.0, rentals=600.0, miles=60, margin=0.42, sc=0.10)),
])
EJ_NOTAS = OrderedDict([
    ('Corporate drop-off, 30 guests', 'Chicken & seafood paella for an office, cooked at the commissary and delivered hot '
                                      'in its pan (20-guest minimum). It lands within 10% of the $22 drop-off line of the '
                                      'example mix in the Word plan (section 3).'),
    ('Backyard paella party, 45 guests', 'Seafood paella cooked live in two 36-in pans in a backyard, with salad and bread '
                                         '(20-guest minimum). Within 10% of the $34 backyard party line of the example mix '
                                         'in the Word plan, inside the $12-31 per person of live paella plus the paella '
                                         'chef that other operators charge separately (cafesevilla.com).'),
    ('Wedding, 110 guests, tapas hour + full staff', 'Wedding with a tapas hour, two paella chefs, a helper and five '
                                                     'servers, dishes and flatware rented and passed through, and a 10% '
                                                     'service charge; 50-guest minimum. The price per guest lands within '
                                                     '10% of the $52 wedding line of the example mix in the Word plan, '
                                                     'inside the $44-65 of a full Spanish menu with staff '
                                                     '(miamigrillcatering.com, zola.com).'),
])

# --------------------------------------------------------------------------
# 3. Shopping List by Guest Count: parámetros (Kit estimate) y columnas (tu evento + 40 / 80 / 150 / 300)
# --------------------------------------------------------------------------
COMPRA = [
    ('lemons', 'Lemons per guest (wedges for the pan)', DEC2, 0.25, 'Kit estimate: one lemon cut in 4 wedges.'),
    ('bread', 'Bread rolls or slices per guest', DEC1, 1.5, 'Kit estimate.'),
    ('salad_oz', 'Salad greens per guest (oz)', DEC1, 1.5, 'Kit estimate.'),
    ('aioli_oz', 'Garlic aioli per guest (oz)', DEC1, 0.75, 'Kit estimate: contains egg (allergen).'),
    ('plates', 'Plates and forks per guest', DEC2, 1.1, 'Kit estimate: 10% extra.'),
    ('ice_ratio', 'Ice per lb of raw seafood (lb)', DEC1, 1.0, 'Kit estimate: keep raw seafood at 41 °F (5 °C) or below.'),
    ('cooler_lb', 'Seafood and ice per cooler (lb)', ENTERO, 40, 'Kit estimate: check your cooler size.'),
    ('chafing', 'Guests per chafing dish for the tapas hour', ENTERO, 25, 'Kit estimate. Used only with a tapas hour.'),
]
REF_INVITADOS = (40, 80, 150, 300)
C_COLS = 'BCDEF'


# --------------------------------------------------------------------------
# Lógica en Python (la misma que las fórmulas): la usa el autotest como valor esperado
# --------------------------------------------------------------------------
def _ceil(x):
    return math.ceil(round(x, 9))


def _num(x):
    return x if isinstance(x, (int, float)) and not isinstance(x, bool) else None


def estilos_calc(precios=None, rice_price=None, rice_oz=3.25):
    """Hoja 2, tabla de estilos: coste con TUS precios y nº de ingredientes sin precio, por estilo."""
    precios = precios or [None] * len(INGREDIENTES)
    por_unidad = [(_num(p) or 0) / (16.0 if INGREDIENTES[j][1] == 'lb' else 1.0) for j, p in enumerate(precios)]
    falta = [0 if (_num(p) is not None and p > 0) else 1 for p in precios]
    falta_arroz = 0 if (_num(rice_price) is not None and rice_price > 0) else 1
    out = OrderedDict()
    for nombre, qty, kit, _fc, _al in ESTILOS:
        tuyo = (_num(rice_price) or 0) * rice_oz / 16.0 + sum(q * u for q, u in zip(qty, por_unidad))
        faltan = falta_arroz + sum(f for q, f in zip(qty, falta) if q > 0)
        out[nombre] = OrderedDict([('kit', kit), ('tuyo', tuyo), ('faltan', faltan)])
    return out, por_unidad, falta


def tamanos_calc(rice_oz=3.25, stock_qt=1.2):
    out = OrderedDict()
    for nombre, cm, inv, aros, prop, _n in TAMANOS:
        rice = inv * rice_oz / 16.0
        stock = rice * stock_qt / 4.0
        out[nombre] = OrderedDict([('rice', rice), ('rice_kg', rice * LB_KG), ('stock', stock), ('stock_l', stock * GAL_L),
                                   ('propane_kg', prop * LB_KG)])
    return out


def calcular(e, precios=None, rice_price=None):
    est, _pu, _fa = estilos_calc(precios, rice_price, e['rice_oz'])
    tam = {t[0]: t for t in TAMANOS}[e['pan_size']]
    r = OrderedDict()
    r['billable'] = max(e['guests'], e['minimum'])
    r['guests_pan'] = tam[2]
    r['pans'] = _ceil(e['guests'] / tam[2])
    r['rice_lb'] = e['guests'] * e['rice_oz'] / 16.0
    r['stock_gal'] = r['rice_lb'] * 1.2 / 4.0
    r['propane_lb'] = r['pans'] * tam[4]
    r['cylinders'] = _ceil(r['propane_lb'] / 20.0) + 1
    x = est[e['style']]
    r['missing'] = x['faltan']
    r['food_status'] = ('Complete: food cost uses your prices' if r['missing'] == 0 else
                        'Missing: food cost uses the style estimate per guest')
    r['style_cost'] = x['tuyo'] if r['missing'] == 0 else x['kit']
    r['c_food'] = e['guests'] * (r['style_cost'] + e['salad_bread'])
    r['c_tapas'] = e['guests'] * e['tapas_cost'] if e['tapas'] == 'Yes' else 0
    r['c_disp'] = e['guests'] * e['disposables']
    r['c_propane'] = r['propane_lb'] * e['propane_price']
    r['c_staff'] = (((e['chefs'] * e['wage_chef'] + e['helpers'] * e['wage_helper'] + e['servers'] * e['wage_server'])
                     * e['event_hours'] + e['prep_hours'] * e['wage_helper']) * (1 + e['payroll']))
    r['c_other'] = e['rentals'] + e['miles'] * e['per_mile']
    r['c_total'] = r['c_food'] + r['c_tapas'] + r['c_disp'] + r['c_propane'] + r['c_staff'] + r['c_other']
    r['c_guest'] = r['c_total'] / r['billable']
    r['price'] = _ceil(r['c_total'] / (1 - e['margin'] - e['commission']) / r['billable'] * 2) / 2
    r['subtotal'] = r['price'] * r['billable'] + e['event_fee']
    r['sc_amt'] = r['subtotal'] * e['sc']
    r['tax_amt'] = (r['subtotal'] + (r['sc_amt'] if e['tax_sc'] == 'Yes' else 0)) * e['tax']
    r['total'] = r['subtotal'] + r['sc_amt'] + r['tax_amt']
    r['dep_amt'] = r['total'] * e['deposit']
    r['balance'] = r['total'] - r['dep_amt']
    r['comm_amt'] = r['subtotal'] * e['commission']
    r['profit'] = r['subtotal'] - r['comm_amt'] - r['c_total']
    r['margin_calc'] = 0 if r['subtotal'] == 0 else r['profit'] / r['subtotal']
    r['chefs_sug'] = _ceil(r['pans'] / 3.0)
    r['helper_sug'] = 1 if (r['pans'] >= 4 or e['guests'] >= 80) else 0
    srv = PAQUETES[e['package']][0]
    r['srv_sug'] = _ceil(e['guests'] / srv) if srv > 0 else 0
    r['tiers_check'] = 'OK' if (e['d1'] > e['d2'] > e['d3']) else 'Adjust: tier 1 > tier 2 > tier 3 days'
    for i in range(1, 5):
        r['ret%d' % i] = r['total'] * e['r%d' % i]
    return r


def calcular_compra(e, c=None, ref=REF_INVITADOS):
    """Hoja 3: para los invitados de la hoja 1 y las 4 columnas de referencia (mismo estilo y misma paellera)."""
    c = c or OrderedDict((k, v) for k, _r, _f, v, _n in COMPRA)
    tam = {t[0]: t for t in TAMANOS}[e['pan_size']]
    qty = {x[0]: x[1] for x in ESTILOS}[e['style']]
    srv = PAQUETES[e['package']][0]
    cols = []
    for g in (e['guests'],) + tuple(ref):
        r = OrderedDict()
        r['guests'] = g
        r['pans'] = _ceil(g / tam[2])
        r['rice'] = g * e['rice_oz'] / 16.0
        r['stock'] = r['rice'] * 1.2 / 4.0
        for j, (nombre, uni) in enumerate(INGREDIENTES):
            r['ing%d' % j] = g * qty[j] / (16.0 if uni == 'lb' else 1.0)
        mar = sum(r['ing%d' % j] for j in SEAFOOD_IDX)
        r['seafood'] = mar
        r['lemons'] = _ceil(g * c['lemons'])
        r['bread'] = _ceil(g * c['bread'])
        r['salad'] = g * c['salad_oz'] / 16.0
        r['aioli'] = g * c['aioli_oz'] / 32.0
        r['propane'] = r['pans'] * tam[4]
        r['cyl'] = _ceil(r['propane'] / 20.0) + 1
        r['plates'] = _ceil(g * c['plates'])
        r['ice'] = mar * c['ice_ratio']
        r['coolers'] = _ceil((mar + r['ice']) / c['cooler_lb']) if c['cooler_lb'] > 0 else 0
        r['chafing'] = _ceil(g / c['chafing']) if (e['tapas'] == 'Yes' and c['chafing'] > 0) else 0
        r['chefs'] = _ceil(r['pans'] / 3.0)
        r['helper'] = 1 if (r['pans'] >= 4 or g >= 80) else 0
        r['srv'] = _ceil(g / srv) if srv > 0 else 0
        cols.append(r)
    return cols


# --------------------------------------------------------------------------
# Fórmulas
# --------------------------------------------------------------------------
def a_excel(expr, ref):
    """Plantilla Python-ish → fórmula Excel con las celdas de `ref` (clave → coordenada)."""
    f = (expr.replace('max(', 'MAX(').replace('roundup(', 'ROUNDUP(').replace('if(', 'IF(').replace('or(', 'OR(')
         .replace('==', '='))
    f = re.sub(r'\{(\w+)\}', lambda m: ref[m.group(1)], f)
    f = f.replace('*2)/2', '*2,0)/2')
    f = re.sub(r'ROUNDUP\(([^()]*)\)(?!/2)', r'ROUNDUP(\1,0)', f)
    return f


def _idx(rango, clave_ref, nombres):
    return 'INDEX(%s,MATCH(%s,%s,0))' % (rango, clave_ref, nombres)


def formula_salida(k, expr, ref):
    if expr.startswith('paq:'):
        v = _idx(_TP('B'), ref['package'], T_NOMBRES)
        return '=IFERROR(IF(%s>0,ROUNDUP(%s/%s,0),0),0)' % (v, ref['guests'], v)
    if expr.startswith('tam:'):
        return '=IFERROR(%s,1)' % _idx(_tcol(expr[4:]), ref['pan_size'], T_TAM)
    if expr.startswith('est:'):
        return '=IFERROR(%s,0)' % _idx(_ecol(expr[4:]), ref['style'], T_EST)
    if expr == 'caldo':
        return "=%s*'%s'!$B$%d/4" % (ref['rice_lb'], H2, Y_STOCK)
    if expr == 'propano':
        return '=%s*IFERROR(%s,0)' % (ref['pans'], _idx(_tcol('I'), ref['pan_size'], T_TAM))
    if expr == 'estado':
        return ('=IF(%s=0,"Complete: food cost uses your prices","Missing: food cost uses the style estimate per '
                'guest")' % ref['missing'])
    if expr == 'coste_estilo':
        return '=IFERROR(IF(%s=0,%s,%s),0)' % (ref['missing'], _idx(_ecol('N'), ref['style'], T_EST),
                                               _idx(_ecol('M'), ref['style'], T_EST))
    if expr == 'tramos':
        return '=IF(AND(%s>%s,%s>%s),"OK","Adjust: tier 1 > tier 2 > tier 3 days")' % (ref['d1'], ref['d2'], ref['d2'],
                                                                                      ref['d3'])
    return '=' + a_excel(expr, ref)


# --------------------------------------------------------------------------
# Construcción
# --------------------------------------------------------------------------
def _f(ws, coord, v, b=False, sz=11, i=False, fmt=None, wrap=False, fill=None, locked=True, color=None):
    c = ws[coord]
    c.value = v
    c.font = Font(name=F, size=sz, bold=b, italic=i, color=color)
    if fmt:
        c.number_format = fmt
    c.alignment = Alignment(wrap_text=wrap, vertical='top')
    if fill:
        c.fill = fill
    c.protection = Protection(locked=locked)
    return c


def _imprimir(ws, horizontal=False):
    ws.page_setup.paperSize = M.PAPEL_EN                                         # 1 = US Letter
    ws.page_setup.orientation = 'landscape' if horizontal else 'portrait'
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.oddFooter.center.text = M.FOOTER_EN
    ws.print_options.horizontalCentered = True
    ws.protection.sheet = True                                                   # sin contraseña (D20)
    ws.sheet_view.showGridLines = False


def _dv(ws, tipo, celdas, f1=None, f2=None, op=None, titulo='Invalid value', msg='', blanco=False):
    dv = DataValidation(type=tipo, formula1=f1, formula2=f2, operator=op, allow_blank=blanco, showErrorMessage=True,
                        errorTitle=titulo, error=msg)
    ws.add_data_validation(dv)
    for c in celdas:
        dv.add(c)


def _cabecera(ws, titulo, intro, ancho_merge, alto=48):
    _f(ws, 'A1', titulo, b=True, sz=14)
    _f(ws, 'A2', intro, wrap=True)
    ws.merge_cells('A2:%s2' % ancho_merge)
    ws.row_dimensions[2].height = alto


def _fmt_entrada(fmt):
    return fmt if fmt != '@' else None


def hoja_calculadora(wb):
    ws = wb.active
    ws.title = H1
    for col, w in (('A', 62), ('B', 34), ('C', 80), ('D', 16)):
        ws.column_dimensions[col].width = w
    _cabecera(ws, 'PAELLA CATERING PRICING CALCULATOR — PANS, RICE AND PRICE PER GUEST TO QUOTE',
              'Type in the GREEN cells; everything else is calculated. Pick the package, the menu style and the pan size, '
              'your staff and your costs: the tab turns them into pans to cook, rice, stock and propane to bring, a price '
              'per guest and a full quote. Ingredient prices are typed on the "%s" tab (left blank on purpose). Amounts '
              'are in US dollars (or your own currency) and exclude sales tax unless the row says otherwise. Every '
              'default value is a kit estimate: replace it with your own numbers.' % H2, 'C')
    for col, t in zip('ABC', ('INPUTS', 'Value', 'Notes')):
        _f(ws, '%s4' % col, t, b=True, fill=GRIS)
    for k, rot, fmt, val, nota in ENTRADAS:
        f = K[k]
        _f(ws, 'A%d' % f, rot)
        _f(ws, 'B%d' % f, val, fmt=_fmt_entrada(fmt), fill=VERDE, locked=False, wrap=k in ('package', 'style'))
        if nota:
            _f(ws, 'C%d' % f, nota, wrap=True)
        if len(nota) > 90 or k == 'package':
            ws.row_dimensions[f].height = 15 * math.ceil(len(nota) / 90.0) + 2
    for col, t in zip('ABC', ('CALCULATION', 'Value', 'Notes')):
        _f(ws, '%s%d' % (col, F_CALC), t, b=True, fill=GRIS)
    ref = {k: '$B$%d' % f for k, f in K.items()}
    ref.update({k: '$B$%d' % f for k, f in S.items()})
    for k, rot, fmt, expr, nota in SALIDAS:
        f = S[k]
        neg = k in ('price', 'total', 'margin_calc', 'food_status')
        _f(ws, 'A%d' % f, rot, b=neg)
        _f(ws, 'B%d' % f, formula_salida(k, expr, ref), fmt=_fmt_entrada(fmt), b=neg, wrap=fmt == '@')
        if nota:
            _f(ws, 'C%d' % f, nota, wrap=True)
            if len(nota) > 90:
                ws.row_dimensions[f].height = 15 * math.ceil(len(nota) / 90.0) + 2
        elif fmt == '@':
            ws.row_dimensions[f].height = 32
    cab = ('PACKAGES (Kit estimate — edit the numbers)', 'Guests per server or attendant (0 = none)', 'Notes')
    for col, t in zip('ABC', cab):
        _f(ws, '%s%d' % (col, F_TABLA), t, b=True, fill=GRIS, wrap=True)
    ws.row_dimensions[F_TABLA].height = 32
    for i, (paq, (srv, nota)) in enumerate(PAQUETES.items(), 1):
        f = F_TABLA + i
        _f(ws, 'A%d' % f, paq, wrap=True)
        _f(ws, 'B%d' % f, srv, fmt=ENTERO, fill=VERDE, locked=False)
        _f(ws, 'C%d' % f, nota, wrap=True)
        ws.row_dimensions[f].height = 48
    # validaciones (D123)
    _dv(ws, 'list', [ref['package'].replace('$', '')], f1=T_NOMBRES.replace("'%s'!" % H1, ''),
        titulo='Invalid package', msg='Choose a package from the list.')
    _dv(ws, 'list', [ref['style'].replace('$', '')], f1=T_EST, titulo='Invalid style',
        msg='Choose a menu style from the list on the "%s" tab.' % H2)
    _dv(ws, 'list', [ref['pan_size'].replace('$', '')], f1=T_TAM, titulo='Invalid pan size',
        msg='Choose a pan size from the list on the "%s" tab.' % H2)
    _dv(ws, 'list', [ref[k].replace('$', '') for k in ('tax_sc', 'tapas')], f1='"Yes,No"', titulo='Invalid entry',
        msg='Choose Yes or No.')
    _dv(ws, 'whole', ['B%d' % K['guests']], f1='1', f2='100000', op='between', titulo='Invalid count',
        msg='Enter a whole number of 1 or more.')
    _dv(ws, 'whole', ['B%d' % K[k] for k in ('minimum', 'chefs', 'helpers', 'servers', 'miles')], f1='0', f2='100000',
        op='between', titulo='Invalid count', msg='Enter a whole number of 0 or more.')
    _dv(ws, 'whole', ['B%d' % K[k] for k in ('d1', 'd2', 'd3')], f1='1', f2='730', op='between', titulo='Invalid days',
        msg='Enter a whole number of days between 1 and 730.')
    _dv(ws, 'decimal', ['B%d' % K['rice_oz']], f1='1', f2='8', op='between', titulo='Invalid amount',
        msg='Enter the uncooked rice per guest in oz, between 1 and 8.')
    _dv(ws, 'decimal', ['B%d' % K[k] for k in ('tapas_cost', 'salad_bread', 'disposables', 'propane_price', 'event_hours',
                                              'prep_hours', 'wage_chef', 'wage_helper', 'wage_server', 'rentals',
                                              'per_mile', 'event_fee')],
        f1='0', op='greaterThanOrEqual', titulo='Invalid amount', msg='Enter a number greater than or equal to 0.')
    _dv(ws, 'decimal', ['B%d' % K[k] for k in ('payroll', 'sc', 'tax', 'deposit', 'r1', 'r2', 'r3', 'r4')],
        f1='0', f2='1', op='between', titulo='Invalid percentage', msg='Enter a percentage between 0 and 1 (0.20 = 20%).')
    _dv(ws, 'decimal', ['B%d' % K['commission']], f1='0', f2='0.3', op='between', titulo='Invalid percentage',
        msg='Enter a fee between 0 and 30% (0.15 = 15%).')
    _dv(ws, 'decimal', ['B%d' % K['margin']], f1='0', f2='0.65', op='between', titulo='Invalid percentage',
        msg='Enter a target margin between 0 and 65% (0.45 = 45%).')
    _dv(ws, 'whole', ['B%d' % (F_TABLA + i) for i in range(1, len(PAQUETES) + 1)], f1='0', f2='10000', op='between',
        titulo='Invalid count', msg='Enter a whole number of guests, 0 or more.')
    ws.freeze_panes = 'A5'
    _imprimir(ws)
    return ws


def hoja_yield(wb):
    ws = wb.create_sheet(H2)
    anchos = (('A', 40), ('B', 12), ('C', 12), ('D', 12), ('E', 12), ('F', 12), ('G', 12), ('H', 12), ('I', 12),
              ('J', 12), ('K', 14), ('L', 10), ('M', 12), ('N', 12), ('O', 12), ('P', 34), ('Q', 30), ('R', 40))
    for col, w in anchos:
        ws.column_dimensions[col].width = w
    _cabecera(ws, 'PAELLA YIELD & PAN SIZING — 7 PAN SIZES AND 8 MENU STYLES',
              'Guests, rice, stock and propane per pan size, and the ingredients per guest of 8 menu styles (green cells: '
              'kit estimates, edit them with your own test cooks). Type your ingredient prices in the green row under the '
              'styles: they are blank on purpose. The FDA Food Code MINIMUM cooking temperature (the legal safety floor, '
              '2022 Food Code 3-401.11) and the hot holding rule sit in their own columns.', 'R', alto=60)
    _f(ws, 'A4', 'SIZING INPUT', b=True, fill=GRIS)
    _f(ws, 'A%d' % Y_STOCK, 'Stock per lb of uncooked rice (quarts)')
    _f(ws, 'B%d' % Y_STOCK, 1.2, fmt=DEC2, fill=VERDE, locked=False)
    _f(ws, 'C%d' % Y_STOCK, 'Kit estimate: 1.2 quarts per lb (2.5 liters per kg); bomba rice absorbs up to 3 times its '
                            'volume. The uncooked rice per guest comes from the "%s" tab.' % H1)
    cab = ('Pan size (in, cm)', 'Diameter (cm)', 'Guests per pan (Kit estimate)', 'Burner rings', 'Uncooked rice per pan '
           '(lb)', '(kg)', 'Stock per pan (gallons)', '(liters)', 'Propane per cook (lb)', '(kg)', 'Notes')
    fila = Y_A0 - 1
    for col, t in zip('ABCDEFGHIJK', cab):
        _f(ws, '%s%d' % (col, fila), t, b=True, fill=GRIS, wrap=True)
    ws.row_dimensions[fila].height = 45
    rice_oz = "'%s'!$B$%d" % (H1, K['rice_oz'])
    verdes = []
    for i, (nombre, cm, inv, aros, prop, nota) in enumerate(TAMANOS):
        r = Y_A0 + i
        _f(ws, 'A%d' % r, nombre)
        _f(ws, 'B%d' % r, cm, fmt=ENTERO)
        _f(ws, 'C%d' % r, inv, fmt=ENTERO, fill=VERDE, locked=False)
        _f(ws, 'D%d' % r, aros, fmt=ENTERO, fill=VERDE, locked=False)
        _f(ws, 'E%d' % r, '=C%d*%s/16' % (r, rice_oz), fmt=DEC1)
        _f(ws, 'F%d' % r, '=E%d*%s' % (r, LB_KG), fmt=DEC1)
        _f(ws, 'G%d' % r, '=E%d*$B$%d/4' % (r, Y_STOCK), fmt=DEC2)
        _f(ws, 'H%d' % r, '=G%d*%s' % (r, GAL_L), fmt=DEC1)
        _f(ws, 'I%d' % r, prop, fmt=DEC1, fill=VERDE, locked=False)
        _f(ws, 'J%d' % r, '=I%d*%s' % (r, LB_KG), fmt=DEC1)
        if nota:
            _f(ws, 'K%d' % r, nota, wrap=True)
        ws.row_dimensions[r].height = 30
        verdes += ['C%d' % r, 'I%d' % r]
    _dv(ws, 'whole', ['C%d' % (Y_A0 + i) for i in range(len(TAMANOS))], f1='1', f2='500', op='between',
        titulo='Invalid count', msg='Enter a whole number of guests per pan, 1 or more.')
    _dv(ws, 'whole', ['D%d' % (Y_A0 + i) for i in range(len(TAMANOS))], f1='1', f2='4', op='between',
        titulo='Invalid count', msg='Enter the burner rings, 1 to 4.')
    _dv(ws, 'decimal', ['I%d' % (Y_A0 + i) for i in range(len(TAMANOS))], f1='0', op='greaterThanOrEqual',
        titulo='Invalid amount', msg='Enter a number greater than or equal to 0.')
    _dv(ws, 'decimal', ['B%d' % Y_STOCK], f1='0.5', f2='4', op='between', titulo='Invalid amount',
        msg='Enter the quarts of stock per lb of rice, between 0.5 and 4.')
    # tabla de estilos
    cab = ['Menu style'] + [x[0] for x in INGREDIENTES] + ['Kit food cost per guest (excl. salad & bread)',
                                                          'Food cost per guest with your prices',
                                                          'Ingredients with no price (rice included)',
                                                          'FDA Food Code minimum (safety)', 'Hot holding (Food Code)',
                                                          'Typical allergens']
    for col, t in zip('ABCDEFGHIJKLMNOPQR', cab):
        _f(ws, '%s%d' % (col, Y_B_CAB), t, b=True, fill=GRIS, wrap=True)
    ws.row_dimensions[Y_B_CAB].height = 75
    _f(ws, 'A%d' % (Y_B_CAB - 1), 'MENU STYLES — ingredients per guest (oz; saffron in grams), Kit estimate', b=True)
    arroz = '$B$%d' % Y_ARROZ
    for i, (nombre, qty, kit, fc, alerg) in enumerate(ESTILOS):
        r = Y_B0 + i
        _f(ws, 'A%d' % r, nombre, wrap=True)
        for j, q in enumerate(qty):
            _f(ws, '%s%d' % (ING_COLS[j], r), q, fmt=DEC2, fill=VERDE, locked=False)
        _f(ws, 'M%d' % r, kit, fmt=MONEDA, fill=VERDE, locked=False)
        _f(ws, 'N%d' % r, '=IF(ISNUMBER(%s),%s,0)*%s/16+SUMPRODUCT(B%d:L%d,$B$%d:$L$%d)' % (
            arroz, arroz, rice_oz, r, r, Y_HELP_P, Y_HELP_P), fmt=MONEDA)
        terminos = ['IF(%s%d>0,%s$%d,0)' % (c, r, c, Y_HELP_F) for c in ING_COLS]
        _f(ws, 'O%d' % r, '=IF(ISNUMBER(%s),IF(%s>0,0,1),1)+%s' % (arroz, arroz, '+'.join(terminos)), fmt=ENTERO)
        _f(ws, 'P%d' % r, fc, wrap=True)
        _f(ws, 'Q%d' % r, _HH, wrap=True)
        _f(ws, 'R%d' % r, alerg, wrap=True)
        ws.row_dimensions[r].height = 62
        verdes += ['%s%d' % (c, r) for c in ING_COLS + 'M']
    _f(ws, 'A%d' % Y_PRECIO, 'YOUR PRICES: per lb (saffron: per gram). Blank on purpose: type your supplier prices',
       b=True, wrap=True)
    for j, (nombre, uni) in enumerate(INGREDIENTES):
        c = ING_COLS[j]
        _f(ws, '%s%d' % (c, Y_PRECIO), None, fmt=MONEDA, fill=VERDE, locked=False)
        div = '/16' if uni == 'lb' else ''
        _f(ws, '%s%d' % (c, Y_HELP_P), '=IF(ISNUMBER(%s%d),%s%d%s,0)' % (c, Y_PRECIO, c, Y_PRECIO, div), fmt='0.000',
           i=True)
        _f(ws, '%s%d' % (c, Y_HELP_F), '=IF(ISNUMBER(%s%d),IF(%s%d>0,0,1),1)' % (c, Y_PRECIO, c, Y_PRECIO), fmt=ENTERO,
           i=True)
    ws.row_dimensions[Y_PRECIO].height = 45
    _f(ws, 'A%d' % Y_HELP_P, 'Calculated: your price per oz (saffron: per gram)', i=True)
    _f(ws, 'A%d' % Y_HELP_F, 'Calculated: price missing? (1 = yes)', i=True)
    _f(ws, 'A%d' % Y_ARROZ, 'YOUR PRICE: bomba rice per lb (blank on purpose)', b=True)
    _f(ws, 'B%d' % Y_ARROZ, None, fmt=MONEDA, fill=VERDE, locked=False)
    _f(ws, 'C%d' % Y_ARROZ, 'Bomba or another short-grain Spanish rice, from an importer or a cash-and-carry.')
    _dv(ws, 'decimal', ['%s%d' % (c, Y_PRECIO) for c in ING_COLS] + ['B%d' % Y_ARROZ], f1='0', op='greaterThanOrEqual',
        titulo='Invalid price', msg='Type your price (a number of 0 or more), or leave it blank.', blanco=True)
    _dv(ws, 'decimal', verdes, f1='0', op='greaterThanOrEqual', titulo='Invalid amount',
        msg='Enter a number greater than or equal to 0.')
    fila = Y_ARROZ + 2
    for t in ('Food Code minimum = the legal safety floor for each ingredient (3-401.11): chicken 165 °F (74 °C); fish and '
              'shellfish, crustaceans and mollusks included, 145 °F (63 °C) for 15 seconds; raw Mexican-style chorizo is '
              'ground meat, 155 °F (68 °C) for 17 seconds. Spanish cured chorizo is ready to eat. Probe the thickest '
              'piece of chicken in every pan.',
              'Cooked rice needs time and temperature control: Bacillus cereus is common in outbreaks from poor hot '
              'holding of starchy foods like rice (Food Code, Annex 3). Serve at 135 °F (57 °C) or above (3-501.16), or '
              'use time as a public health control with a written procedure: mark the time, discard after 4 hours '
              '(3-501.19), if your county accepts it.',
              'Leftover paella goes in the trash: do not cool it down to serve it again.',
              'Mussels and clams: buy them from an approved source with the dealer tag, keep them in the container until '
              'it is empty, write the date of the last one served and keep the tag for 90 days (3-203.12). Keep raw '
              'seafood at 41 °F (5 °C) or below.',
              'Allergens: the 9 major US allergens include fish and crustacean shellfish (crab, lobster, shrimp); '
              'mollusks are not one of them, but declare them anyway. Garlic aioli has egg; bread and fideua (noodle '
              'paella) have wheat. In the UK, mollusks, celery and sulphites are among the 14 allergens. The column above '
              'lists what is typical: verify with your recipe and suppliers.',
              'Guests per pan, propane and ingredients per guest are kit estimates: weigh your own rice, stock and '
              'propane on your test cooks and type them here; the "%s" and "%s" tabs follow.' % (H1, H3)):
        _f(ws, 'A%d' % fila, t, wrap=True)
        ws.merge_cells('A%d:R%d' % (fila, fila))
        ws.row_dimensions[fila].height = 32
        fila += 1
    ws.freeze_panes = 'B%d' % Y_A0
    _imprimir(ws, horizontal=True)
    return ws


def hoja_compra(wb):
    ws = wb.create_sheet(H3)
    for col, w in (('A', 56), ('B', 16), ('C', 14), ('D', 14), ('E', 14), ('F', 14), ('G', 56)):
        ws.column_dimensions[col].width = w
    _cabecera(ws, 'SHOPPING LIST BY GUEST COUNT — WHAT TO BUY AND BRING',
              'It reads the package, the menu style, the pan size, the rice per guest and the tapas hour switch from the '
              '"%s" tab and the pans and ingredients from the "%s" tab. Column B is your event; columns C-F are reference '
              'sizes (40, 80, 150 and 300 guests: edit them). Green cells are kit estimates: edit them.' % (H1, H2), 'G')
    g1 = "'%s'!" % H1
    R = {}
    fila = 4
    _f(ws, 'A%d' % fila, 'ASSUMPTIONS (Kit estimate — edit)', b=True, fill=GRIS)
    for col, t in zip('BG', ('Value', 'Notes')):
        _f(ws, '%s%d' % (col, fila), t, b=True, fill=GRIS)
    for k, rot, fmt, val, nota in COMPRA:
        fila += 1
        _f(ws, 'A%d' % fila, rot)
        _f(ws, 'B%d' % fila, val, fmt=fmt, fill=VERDE, locked=False)
        if nota:
            _f(ws, 'G%d' % fila, nota, wrap=True)
        R[k] = '$B$%d' % fila
    fila += 2
    f_cab = fila
    _f(ws, 'A%d' % fila, 'QUANTITIES', b=True, fill=GRIS)
    _f(ws, 'B%d' % fila, 'Your event', b=True, fill=GRIS)
    for j in range(4):
        _f(ws, '%s%d' % (C_COLS[j + 1], fila), 'Reference', b=True, fill=GRIS)
    _f(ws, 'G%d' % fila, 'Notes', b=True, fill=GRIS)
    inv_pan = 'IFERROR(%s,1)' % _idx(_tcol('C'), '%s$B$%d' % (g1, K['pan_size']), T_TAM)
    prop_pan = 'IFERROR(%s,0)' % _idx(_tcol('I'), '%s$B$%d' % (g1, K['pan_size']), T_TAM)
    srv = 'IFERROR(%s,0)' % _idx(_TP('B'), '%s$B$%d' % (g1, K['package']), T_NOMBRES)
    rice_oz = '%s$B$%d' % (g1, K['rice_oz'])
    tapas = '%s$B$%d' % (g1, K['tapas'])
    filas = OrderedDict()

    def linea(clave, rot, fx, fmt, nota=''):
        nonlocal fila
        fila += 1
        _f(ws, 'A%d' % fila, rot)
        for col in C_COLS:
            _f(ws, '%s%d' % (col, fila), fx(col), fmt=fmt)
        if nota:
            _f(ws, 'G%d' % fila, nota, wrap=True)
        filas[clave] = fila

    fila += 1
    _f(ws, 'A%d' % fila, 'Guests', b=True)
    _f(ws, 'B%d' % fila, '=%s$B$%d' % (g1, K['guests']), fmt=ENTERO, b=True)
    for j, n in enumerate(REF_INVITADOS):
        _f(ws, '%s%d' % (C_COLS[j + 1], fila), n, fmt=ENTERO, b=True, fill=VERDE, locked=False)
    _f(ws, 'G%d' % fila, 'Column B comes from the "%s" tab.' % H1, wrap=True)
    filas['guests'] = fila
    gr = lambda col: '%s$%d' % (col, filas['guests'])                           # noqa: E731
    linea('pans', '="Paella pans: "&%s$B$%d' % (g1, K['pan_size']), lambda c: '=ROUNDUP(%s/%s,0)' % (gr(c), inv_pan),
          ENTERO, 'One 3-ring or 2-ring burner per pan (see the pan table).')
    ws['A%d' % filas['pans']].value = '="Paella pans and burners: "&%s$B$%d' % (g1, K['pan_size'])
    linea('rice', 'Uncooked bomba rice (lb)', lambda c: '=%s*%s/16' % (gr(c), rice_oz), DEC1, 'Guests × oz per guest ÷ 16 (1 lb = 0.454 kg).')
    linea('stock', 'Stock (gallons)', lambda c: "=%s$%d*'%s'!$B$%d/4" % (c, filas['rice'], H2, Y_STOCK), DEC1,
          'Make it the day before at the commissary and keep it cold (1 gallon = 3.8 liters).')
    for j, (nombre, uni) in enumerate(INGREDIENTES):
        q = 'INDEX(%s,MATCH(%s$B$%d,%s,0))' % (_ecol(ING_COLS[j]), g1, K['style'], T_EST)
        div = '/16' if uni == 'lb' else ''
        rot = nombre.replace(' (oz)', ' (lb)')
        linea('ing%d' % j, rot, lambda c, q=q, div=div: '=IFERROR(%s*%s%s,0)' % (gr(c), q, div),
              DEC2 if uni == 'g' else DEC1, 'For the menu style on the "%s" tab.' % H1 if j == 0 else '')
    linea('seafood', 'Raw seafood, all lines (lb)', lambda c: '=' + '+'.join('%s$%d' % (c, filas['ing%d' % j])
                                                                             for j in SEAFOOD_IDX), DEC1, '')
    linea('lemons', 'Lemons', lambda c: '=ROUNDUP(%s*%s,0)' % (gr(c), R['lemons']), ENTERO, '')
    linea('bread', 'Bread rolls or slices', lambda c: '=ROUNDUP(%s*%s,0)' % (gr(c), R['bread']), ENTERO, '')
    linea('salad', 'Salad greens (lb)', lambda c: '=%s*%s/16' % (gr(c), R['salad_oz']), DEC1, '')
    linea('aioli', 'Garlic aioli (quarts)', lambda c: '=%s*%s/32' % (gr(c), R['aioli_oz']), DEC1, 'Egg allergen (1 quart = 0.95 liters).')
    linea('propane', 'Propane (lb)', lambda c: '=%s$%d*%s' % (c, filas['pans'], prop_pan), DEC1, '')
    linea('cyl', '20-lb propane cylinders (plus one spare)', lambda c: '=ROUNDUP(%s$%d/20,0)+1' % (c, filas['propane']),
          ENTERO, 'They travel upright and secured in the van.')
    linea('plates', 'Plates and forks', lambda c: '=ROUNDUP(%s*%s,0)' % (gr(c), R['plates']), ENTERO, '')
    linea('ice', 'Ice for the raw seafood (lb)', lambda c: '=%s$%d*%s' % (c, filas['seafood'], R['ice_ratio']), DEC1, '')
    linea('coolers', 'Coolers for the raw seafood', lambda c: '=IF(%s>0,ROUNDUP((%s$%d+%s$%d)/%s,0),0)' % (
        R['cooler_lb'], c, filas['seafood'], c, filas['ice'], R['cooler_lb']), ENTERO, '')
    linea('chafing', 'Chafing dishes for the tapas hour', lambda c: '=IF(AND(%s="Yes",%s>0),ROUNDUP(%s/%s,0),0)' % (
        tapas, R['chafing'], gr(c), R['chafing']), ENTERO, '0 without a tapas hour.')
    linea('chefs', 'Paella chefs (suggested: one per 3 pans)', lambda c: '=ROUNDUP(%s$%d/3,0)' % (c, filas['pans']),
          ENTERO, 'Kit estimate.')
    linea('helper', 'Helper (suggested: from 4 pans or 80 guests)', lambda c: '=IF(OR(%s$%d>=4,%s>=80),1,0)' % (
        c, filas['pans'], gr(c)), ENTERO, 'Kit estimate.')
    linea('srv', 'Servers and attendants (suggested)', lambda c: '=IF(%s>0,ROUNDUP(%s/%s,0),0)' % (srv, gr(c), srv),
          ENTERO, 'From the package table of the "%s" tab: 0 for drop-off and live paella parties.' % H1)
    _dv(ws, 'whole', ['%s%d' % (C_COLS[j + 1], filas['guests']) for j in range(4)], f1='1', f2='100000', op='between',
        titulo='Invalid count', msg='Enter a whole number of guests, 1 or more.')
    _dv(ws, 'whole', [R[k].replace('$', '') for k in ('cooler_lb', 'chafing')], f1='0', f2='10000',
        op='between', titulo='Invalid count', msg='Enter a whole number of 0 or more.')
    _dv(ws, 'decimal', [R[k].replace('$', '') for k in ('lemons', 'bread', 'salad_oz', 'aioli_oz', 'plates', 'ice_ratio')],
        f1='0', op='greaterThanOrEqual', titulo='Invalid amount', msg='Enter a number greater than or equal to 0.')
    ws.freeze_panes = 'B%d' % (f_cab + 1)
    _imprimir(ws, horizontal=True)
    return ws, filas


C_CLAVES = (('guests', 'pans', 'rice', 'stock') + tuple('ing%d' % j for j in range(len(INGREDIENTES))) +
            ('seafood', 'lemons', 'bread', 'salad', 'aioli', 'propane', 'cyl', 'plates', 'ice', 'coolers', 'chafing',
             'chefs', 'helper', 'srv'))


def hoja_ejemplos(wb):
    ws = wb.create_sheet(H4)
    ws.column_dimensions['A'].width = 62
    for col in 'BCD':
        ws.column_dimensions[col].width = 30
    _cabecera(ws, '3 EXAMPLE QUOTES — calculated with the same formulas as the "%s" tab' % H1,
              'Kit estimates for illustration. No ingredient price is typed: until you type yours on the "%s" tab, the '
              'food cost uses the style estimate per guest. Each quote lands within 10%% of the price of its line in the '
              'example mix in the Word plan (section 3) ($22 drop-off, $34 backyard party, $52 wedding). Change the '
              'inputs in "%s" to price your own events.' % (H2, H1), 'D')
    _f(ws, 'A4', 'INPUTS', b=True, fill=GRIS)
    for j, nombre in enumerate(EJEMPLOS):
        _f(ws, '%s4' % 'BCD'[j], nombre, b=True, fill=GRIS, wrap=True)
    ws.row_dimensions[4].height = 45
    for k, rot, fmt, val, nota in ENTRADAS:
        _f(ws, 'A%d' % K[k], rot)
        for j, e in enumerate(EJEMPLOS.values()):
            _f(ws, '%s%d' % ('BCD'[j], K[k]), e[k], fmt=_fmt_entrada(fmt), wrap=k in ('package', 'style'))
        if k in ('package', 'style'):
            ws.row_dimensions[K[k]].height = 32
    _f(ws, 'A%d' % F_CALC, 'RESULT', b=True, fill=GRIS)
    for j in range(3):
        _f(ws, '%s%d' % ('BCD'[j], F_CALC), list(EJEMPLOS)[j], b=True, fill=GRIS, wrap=True)
    ws.row_dimensions[F_CALC].height = 45
    for k, rot, fmt, expr, nota in SALIDAS:
        neg = k in ('price', 'total', 'margin_calc', 'food_status')
        _f(ws, 'A%d' % S[k], rot, b=neg)
        for j in range(3):
            col = 'BCD'[j]
            ref = {kk: '$%s$%d' % (col, f) for kk, f in K.items()}
            ref.update({kk: '$%s$%d' % (col, f) for kk, f in S.items()})
            _f(ws, '%s%d' % (col, S[k]), formula_salida(k, expr, ref), fmt=_fmt_entrada(fmt), b=neg, wrap=fmt == '@')
        if fmt == '@':
            ws.row_dimensions[S[k]].height = 32
    f = S[SALIDAS[-1][0]] + 2
    _f(ws, 'A%d' % f, 'Price per guest with the event fee spread over the billable guests', b=True)
    for j in range(3):
        col = 'BCD'[j]
        _f(ws, '%s%d' % (col, f), '=$%s$%d/$%s$%d' % (col, S['subtotal'], col, S['billable']), fmt=MONEDA, b=True)
    f_equiv = f
    f += 2
    _f(ws, 'A%d' % f, 'WHAT EACH EXAMPLE SHOWS', b=True, fill=GRIS)
    for j, (nombre, nota) in enumerate(EJ_NOTAS.items()):
        _f(ws, 'A%d' % (f + 1 + j), nombre, b=True, wrap=True)
        _f(ws, 'B%d' % (f + 1 + j), nota, wrap=True)
        ws.merge_cells('B%d:D%d' % (f + 1 + j, f + 1 + j))
        ws.row_dimensions[f + 1 + j].height = 62
    ws.freeze_panes = 'B5'
    _imprimir(ws, horizontal=True)
    return ws, f_equiv


INSTRUCCIONES = [
    ('HOW TO USE THE PAELLA CATERING PRICING CALCULATOR', True),
    ('This workbook turns guests into paella pans, rice, stock and propane, prices each guest and builds a full quote and '
     'the shopping list. You type only in the GREEN cells; the gross margin is calculated, never typed.', False),
    ('1. "%s": choose the package (drop-off paella, live paella party or full-service live paella), the menu style, the '
     'guests, your guest minimum and the pan size. The pan size sets the guests per pan (36 in (90 cm) = about 40, Kit '
     'estimate). The rice per guest is 3.25 oz (92 g) of uncooked bomba rice by default, inside the 3 to 3.5 oz (85-100 '
     'g) that serves a crowd.' % H1, False),
    ('2. "%s": type your supplier prices for the rice and the ingredients of your styles (per lb; saffron per gram). '
     'They are blank on purpose: this kit does not guess 2026 seafood or saffron prices. While any ingredient of the '
     'chosen style has no price, the food cost uses the style estimate per guest (Kit: about $5 chicken & chorizo, $8.50 '
     'seafood, $16.50 lobster) and the "Ingredient prices" row says so.' % H2, False),
    ('3. Back on the first tab, enter the tapas hour switch and its cost, salad and bread, disposables, propane, your '
     'staff (event day and commissary prep hours), rentals and mileage. Then the marketplace fee, your target margin, '
     'the service charge, the sales tax rate and the switch "Does your state tax the service charge?". The price per '
     'guest is the direct cost divided by (1 - target margin - marketplace fee) and by the billable guests, rounded up '
     'to the next 0.50; any event fee goes on top.', False),
    ('4. Deposit and cancellation: US caterers ask 25-50% to hold the date (the calculator uses 50%, a kit estimate) '
     'and four cancellation tiers with no gaps tell you what you keep. Put the same tiers in your contract, with a plan '
     'B for wind, rain or a venue that does not allow open flame.', False),
    ('5. "%s": guests, rice, stock and propane per pan for 7 sizes, and the ingredients per guest, the Food Code minimum, '
     'the hot holding rule and the typical allergens of 8 menu styles. Replace the kit estimates with the numbers from '
     'your own test cooks.' % H2, False),
    ('6. "%s": pans and burners, rice, stock, every ingredient of the style, saffron, lemons, bread, salad, aioli, '
     'propane and cylinders, plates, ice and coolers for the seafood, chafing dishes for the tapas hour and the '
     'suggested team, for your event and for 40, 80, 150 and 300 guests.' % H3, False),
    ('7. "%s": three complete quotes calculated with the same formulas. To move from these per-package prices to the '
     'average price per guest of the financial projections, weight each package price by its guests in the year.' % H4,
     False),
    ('FOOD SAFETY (FDA FOOD CODE 2022)', True),
    ('Cook to the minimum (3-401.11): chicken 165 °F (74 °C); fish and shellfish 145 °F (63 °C) for 15 seconds; raw '
     'Mexican-style chorizo 155 °F (68 °C) for 17 seconds (Spanish cured chorizo is ready to eat). Cooked rice needs '
     'time and temperature control: hold it at 135 °F (57 °C) or above (3-501.16) or use time as a public health '
     'control with a written procedure and discard it after 4 hours (3-501.19, if your county accepts it). Leftovers go '
     'in the trash. Keep the shellstock tags of mussels and clams for 90 days (3-203.12).', False),
    ('SERVICE CHARGE, GRATUITY AND SALES TAX (US)', True),
    ('A mandatory service charge is not a tip: the part you pay out to the staff is wages, with payroll taxes (IRS Rev. '
     'Rul. 2012-18), and the part you keep is revenue. A voluntary gratuity is the client\'s decision and belongs to the '
     'staff: this calculator keeps it out of the price and the total.', False),
    ('Sales tax: many states tax the whole catering invoice (food, staff, rentals and fees); whether the service charge '
     'is taxable depends on the state (California taxes a mandatory one: cdtfa.ca.gov). Set your rate and the switch.',
     False),
    ('COOKING AT THE EVENT AND PROPANE', True),
    ('Cooking the paella at the event is the product, but ask your county first whether your catering permit covers it '
     '(in California, limited food preparation is allowed at private catered events). Ask your fire marshal where the '
     'propane station can go: outdoors or under its own canopy, away from the guest tent, with an extinguisher. Sangria: '
     'without a liquor license, serve it alcohol-free or made with wine the client buys, with no markup.', False),
    ('UK: if you are VAT-registered, enter 0.20 as the tax rate: VAT applies to catering for events and functions.', False),
    ('KIT ESTIMATES', True),
    ('Every default value, the guests per pan, the propane per cook, the ingredients and food cost per guest of each '
     'style and the three example quotes are kit estimates for illustration, not market prices. Replace them with your '
     'own wages, supplier prices, test cooks and tax rate before you quote.', False),
]


def hoja_instrucciones(wb):
    ws = wb.create_sheet(H5)
    ws.column_dimensions['A'].width = 110
    f = 1
    for t, b in INSTRUCCIONES:
        _f(ws, 'A%d' % f, t, b=b, sz=13 if f == 1 else 11, wrap=True)
        if not b:
            ws.row_dimensions[f].height = max(16, 15 * math.ceil(len(t) / 120.0))
        f += 2 if b and f > 1 else 1
    f += 1
    for t in ('To edit a cell that is not green: Review → Unprotect Sheet (there is no password).',
              'More templates and kits in the catalog at ' + M.URL_TIENDA, M.FIRMA_EN, M.version_en('pael'),
              M.AVISO_EN):
        _f(ws, 'A%d' % f, t, wrap=True, i=t == M.AVISO_EN)
        f += 1
    _imprimir(ws)
    return ws


def docprops(wb):
    p = M.PLANES['pael']
    pr = wb.properties
    pr.title = '%s · %s' % (TITULO, p['producto'])
    pr.subject = '%s · v%s' % (p['producto'], M.VERSION)
    pr.keywords = p['keywords']
    pr.description = p['url']
    pr.category = 'AI Chef Pro · Digital products'
    pr.creator = 'AI Chef Pro'
    pr.lastModifiedBy = 'AI Chef Pro'


def generar(destino):
    wb = openpyxl.Workbook()
    hoja_calculadora(wb)
    hoja_yield(wb)
    hoja_compra(wb)
    hoja_ejemplos(wb)
    hoja_instrucciones(wb)
    docprops(wb)
    os.makedirs(os.path.dirname(destino), exist_ok=True)
    wb.save(destino)
    sys.path.insert(0, PD)
    import inject_cache as IC                                                    # noqa: E402
    nf, inj, fail = IC.inject(destino)
    print('%s: formulas=%d cache_inyectado=%d fallos_pycel=%d' % (os.path.relpath(destino, REPO), nf, inj, fail))
    return fail


# --------------------------------------------------------------------------
# Autotest (D124): pycel = caché = Python en las 4 pestañas de cálculo + estilo, tapas hour y precios tecleados probados
# con set_value; ejemplos a ±10 %; ningún precio tecleado
# --------------------------------------------------------------------------
def _cmp(err, etiqueta, got, esperado, tol=0.005):
    if isinstance(esperado, str):
        if got != esperado:
            err.append('%s: %r ≠ esperado %r' % (etiqueta, got, esperado))
        return
    if not isinstance(got, (int, float)) or abs(float(got) - esperado) > tol:
        err.append('%s: %r ≠ esperado %r' % (etiqueta, got, esperado))


PRUEBA_PRECIOS = [5.0, 7.0, 6.5, 9.0, 4.5, 6.0, 18.0, 9.5, 2.0, 1.5, 12.0]       # solo para la prueba
PRUEBA_ARROZ = 3.0


def _escenario(destino, cambios, err, nombre, precios=None, arroz=None):
    """Hoja principal con valores cambiados vía pycel (compilador nuevo: sin estado de otros escenarios)."""
    from pycel import ExcelCompiler
    xl = ExcelCompiler(filename=destino)
    for k in S:                                                     # pycel: la celda tiene que estar en el grafo
        xl.evaluate("'%s'!B%d" % (H1, S[k]))
    e = OrderedDict(_DEF)
    for k, v in cambios.items():
        xl.set_value("'%s'!B%d" % (H1, K[k]), v)
        e[k] = v
    if precios:
        for j, p in enumerate(precios):
            xl.set_value("'%s'!%s%d" % (H2, ING_COLS[j], Y_PRECIO), p)
        xl.set_value("'%s'!B%d" % (H2, Y_ARROZ), arroz)
    esperado = calcular(e, precios, arroz)
    n = 0
    for k, v in esperado.items():
        n += 1
        _cmp(err, '%s · %s!B%d (%s) pycel' % (nombre, H1, S[k], k), xl.evaluate("'%s'!B%d" % (H1, S[k])), v)
    return esperado, n


def autotest(destino):
    from pycel import ExcelCompiler
    err, n = [], 0
    xl = ExcelCompiler(filename=destino)
    wv = openpyxl.load_workbook(destino, data_only=True)
    wf = openpyxl.load_workbook(destino)
    defecto = OrderedDict(_DEF)
    # 1) hoja principal y los 3 ejemplos (pycel + caché)
    casos = [(H4, 'BCD'[j], e, nombre) for j, (nombre, e) in enumerate(EJEMPLOS.items())]
    casos.append((H1, 'B', defecto, 'hoja principal (valores por defecto)'))
    for hoja, col, e, nombre in casos:
        esperado = calcular(e)
        for k, v in esperado.items():
            c = '%s%d' % (col, S[k])
            n += 1
            _cmp(err, '%s · %s!%s (%s) pycel' % (nombre, hoja, c, k), xl.evaluate("'%s'!%s" % (hoja, c)), v)
            _cmp(err, '%s · %s!%s (%s) caché' % (nombre, hoja, c, k), wv[hoja][c].value, v)
        if esperado['margin_calc'] + 1e-9 < e['margin']:
            err.append('%s: margen calculado %.4f < objetivo %.4f' % (nombre, esperado['margin_calc'], e['margin']))
        if esperado['missing'] == 0:
            err.append('%s: precios de ingredientes tecleados (D123: ninguno)' % nombre)
    # 2) los 3 ejemplos a ±10 % del precio de su línea (research §8) — precio por invitado y con la tarifa repartida
    _wsx, f_equiv = hoja_ejemplos(openpyxl.Workbook())
    for j, (nombre, objetivo) in enumerate(OBJETIVO_EJ.items()):
        r = calcular(EJEMPLOS[nombre])
        equiv = r['subtotal'] / r['billable']
        for lab, v in (('precio por invitado', r['price']), ('con tarifas', equiv)):
            if abs(v / objetivo - 1) > 0.10 + 1e-9:
                err.append('%s: %s %.2f fuera de ±10 %% de %.2f' % (nombre, lab, v, objetivo))
        c = '%s%d' % ('BCD'[j], f_equiv)
        n += 1
        _cmp(err, 'equivalente %s pycel' % c, xl.evaluate("'%s'!%s" % (H4, c)), equiv)
        _cmp(err, 'equivalente %s caché' % c, wv[H4][c].value, equiv)
    # 3) Paella Yield & Pan Sizing (pycel + caché): tamaños, estilos y las filas auxiliares
    for i, (nombre, x) in enumerate(tamanos_calc(defecto['rice_oz']).items()):
        r = Y_A0 + i
        for k, v in x.items():
            col = Y_ACOL[k]
            n += 1
            _cmp(err, 'tamaño %s %s%d pycel' % (nombre, col, r), xl.evaluate("'%s'!%s%d" % (H2, col, r)), v, 1e-6)
            _cmp(err, 'tamaño %s %s%d caché' % (nombre, col, r), wv[H2]['%s%d' % (col, r)].value, v, 1e-6)
    est, pu, fa = estilos_calc(None, None, defecto['rice_oz'])
    for i, (nombre, x) in enumerate(est.items()):
        r = Y_B0 + i
        for k in ('tuyo', 'faltan'):
            col = Y_BCOL[k]
            n += 1
            _cmp(err, 'estilo %s %s%d pycel' % (nombre, col, r), xl.evaluate("'%s'!%s%d" % (H2, col, r)), x[k], 1e-6)
            _cmp(err, 'estilo %s %s%d caché' % (nombre, col, r), wv[H2]['%s%d' % (col, r)].value, x[k], 1e-6)
        for col in ('P', 'Q', 'R'):
            if not wf[H2]['%s%d' % (col, r)].value:
                err.append('estilo %s: columna %s vacía' % (nombre, col))
        if '°F' not in (wf[H2]['P%d' % r].value or '') or 'Verify with your recipe' not in (wf[H2]['R%d' % r].value or ''):
            err.append('estilo %s: mínimo del Food Code o aviso de alérgenos ausente' % nombre)
    for j, c in enumerate(ING_COLS):
        for fila, v in ((Y_HELP_P, pu[j]), (Y_HELP_F, fa[j])):
            n += 1
            _cmp(err, 'auxiliar %s%d pycel' % (c, fila), xl.evaluate("'%s'!%s%d" % (H2, c, fila)), v, 1e-9)
            _cmp(err, 'auxiliar %s%d caché' % (c, fila), wv[H2]['%s%d' % (c, fila)].value, v, 1e-9)
    # 4) Shopping List by Guest Count (pycel + caché), con los valores por defecto
    _ws3, filas = hoja_compra(openpyxl.Workbook())
    for j, col in enumerate(C_COLS):
        esperado = calcular_compra(defecto)[j]
        for k in C_CLAVES:
            c = '%s%d' % (col, filas[k])
            n += 1
            _cmp(err, 'compra %s (%s) pycel' % (c, k), xl.evaluate("'%s'!%s" % (H3, c)), esperado[k], 1e-6)
            _cmp(err, 'compra %s (%s) caché' % (c, k), wv[H3][c].value, esperado[k], 1e-6)
    # 5) escenarios en la hoja principal (solo pycel): estilo de langosta, tapas hour, tamaño de paellera, precios
    base = calcular(defecto)
    e_lan, m = _escenario(destino, {'style': 'Lobster & seafood (premium)'}, err, 'estilo langosta')
    n += m
    if not e_lan['price'] > base['price']:
        err.append('estilo de langosta sin efecto')
    e_tap, m = _escenario(destino, {'tapas': 'Yes'}, err, 'tapas hour')
    n += m
    if abs(e_tap['c_tapas'] - defecto['guests'] * defecto['tapas_cost']) > 1e-6 or not e_tap['price'] > base['price']:
        err.append('interruptor «Tapas hour?» sin efecto')
    e_tam, m = _escenario(destino, {'pan_size': '26 in (65 cm)'}, err, 'paellera de 26 in')
    n += m
    if not e_tam['pans'] > base['pans']:
        err.append('tamaño de paellera sin efecto')
    e_pre, m = _escenario(destino, {}, err, 'precios tecleados', PRUEBA_PRECIOS, PRUEBA_ARROZ)
    n += m
    if e_pre['missing'] != 0 or e_pre['food_status'] == base['food_status'] or e_pre['style_cost'] == base['style_cost']:
        err.append('precios tecleados sin efecto')
    for k in ('pans', 'rice_lb', 'stock_gal', 'propane_lb'):                     # las cantidades no dependen del precio
        if abs(e_pre[k] - base[k]) > 1e-9:
            err.append('las cantidades cambian con el precio (%s)' % k)
    # 6) interruptor del sales tax sobre el service charge
    e = dict(EJEMPLOS['Wedding, 110 guests, tapas hour + full staff'])
    if not calcular(dict(e, tax_sc='Yes'))['tax_amt'] > calcular(dict(e, tax_sc='No'))['tax_amt']:
        err.append('interruptor Yes/No del sales tax sin efecto')
    # higiene del fichero (censo-entregables + D20 + D8)
    p = wf.properties
    if p.creator != 'AI Chef Pro' or not (p.title or '').startswith(TITULO) or M.PLANES['pael']['url'] != p.description:
        err.append('docProps D20: %r / %r / %r' % (p.creator, p.title, p.description))
    textos = []
    for ws in wf.worksheets:
        if ws.page_setup.paperSize != 1 or not ws.sheet_properties.pageSetUpPr.fitToPage or \
                ws.oddFooter.center.text != M.FOOTER_EN:
            err.append('%s: impresión (Letter + ajuste + pie)' % ws.title)
        if not ws.protection.sheet:
            err.append('%s: hoja sin proteger' % ws.title)
        for row in ws.iter_rows():
            for c in row:
                v = c.value
                if v == '':
                    err.append('%s!%s: cadena vacía' % (ws.title, c.coordinate))
                if isinstance(v, str) and c.data_type != 'f':
                    textos.append((ws.title, c.coordinate, v))
                if isinstance(v, str) and c.data_type == 'f':
                    textos += [(ws.title, c.coordinate + ' (fórmula)', lit) for lit in re.findall(r'"((?:[^"]|"")*)"', v)
                               if any(ch.isalpha() for ch in lit)]
                if c.data_type == 'f' and wv[ws.title][c.coordinate].value is None:
                    err.append('%s!%s: fórmula sin caché' % (ws.title, c.coordinate))
                if c.fill.fgColor.rgb in ('00E8F5E9', 'FFE8F5E9') and c.protection.locked:
                    err.append('%s!%s: verde bloqueada' % (ws.title, c.coordinate))
                if isinstance(v, str) and c.data_type != 'f' and '€' in v:
                    err.append('%s!%s: «€»' % (ws.title, c.coordinate))
                fmt = c.number_format or ''
                if '€' in fmt or '$' in fmt:
                    err.append('%s!%s: símbolo de moneda en el formato (D8)' % (ws.title, c.coordinate))
        for dv in ws.data_validations.dataValidation:
            for t in (dv.error, dv.errorTitle, dv.prompt, dv.promptTitle):
                if t:
                    textos.append((ws.title, 'DV %s' % dv.sqref, t))
    for campo in ('title', 'subject', 'keywords', 'description'):
        textos.append(('docProps', campo, getattr(p, campo) or ''))
    for h, c, t in textos:
        r = [x for x in M.restos_espanol(t) if not x[0].startswith('oficio')]
        if r:
            err.append('%s!%s: restos de español %s' % (h, c, r[:1]))
        if M.no_latinos(t):
            err.append('%s!%s: caracteres no latinos' % (h, c))
        if M.cifras_prohibidas(t):
            err.append('%s!%s: cifra/término prohibido %s' % (h, c, M.cifras_prohibidas(t)[:1]))
        if re.search(r'(?i)\bvalue\b.*\$\d|\bsave \$|total value|\bwas \$|\d+\s?% off\b', t):
            err.append('%s!%s: ancla de precio' % (h, c))
    # el mínimo del Food Code y el hot holding en columnas propias (D124)
    cab = [wf[H2].cell(Y_B_CAB, j).value for j in range(1, 19)]
    if not any('Food Code minimum' in (x or '') for x in cab) or not any('Hot holding' in (x or '') for x in cab):
        err.append('%s: falta la columna del mínimo del Food Code o la del hot holding' % H2)
    # ningún precio de ingrediente tecleado (D123): las celdas de precio, en blanco
    for c in ['%s%d' % (col, Y_PRECIO) for col in ING_COLS] + ['B%d' % Y_ARROZ]:
        if wf[H2][c].value is not None:
            err.append('%s!%s: precio tecleado' % (H2, c))
    hojas = [ws.title for ws in wf.worksheets]
    if hojas != [H1, H2, H3, H4, H5]:
        err.append('pestañas %r' % hojas)
    for x in err[:80]:
        print('ERROR', x)
    print('autotest calculadora: %d comprobaciones numéricas (pycel + caché) en %s + estilo, tapas hour, paellera y '
          'precios tecleados en la hoja principal · %d textos · %s' % (
              n, ', '.join(HOJAS_CALCULO), len(textos), 'VERDE' if not err else 'ROJO (%d errores)' % len(err)))
    for nombre, e in list(EJEMPLOS.items()) + [('Main tab (defaults)', defecto),
                                               ('Main tab, lobster', dict(defecto, style='Lobster & seafood (premium)')),
                                               ('Main tab, tapas hour', dict(defecto, tapas='Yes'))]:
        r = calcular(e)
        print('  %-48s %3d inv · %d paelleras · %5.1f lb arroz · %4.1f gal caldo · coste %8.2f · precio/inv %6.2f · '
              'subtotal %9.2f · total %9.2f · margen %.1f%% · %s' % (nombre, e['guests'], r['pans'], r['rice_lb'],
                                                                     r['stock_gal'], r['c_total'], r['price'],
                                                                     r['subtotal'], r['total'], r['margin_calc'] * 100,
                                                                     r['food_status'][:8]))
    r = calcular(defecto, PRUEBA_PRECIOS, PRUEBA_ARROZ)
    print('  %-48s coste por invitado del estilo %.2f (Kit %.2f) · precio/inv %.2f' % (
        'Main tab, prices typed (test only)', r['style_cost'], base['style_cost'], r['price']))
    c = calcular_compra(defecto)
    print('  compra: ' + ' · '.join('%d inv: %d paelleras, %.0f lb arroz, %.1f gal caldo, %d cilindros' % (
        x['guests'], x['pans'], x['rice'], x['stock'], x['cyl']) for x in c))
    return 1 if err else 0


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--autotest', action='store_true', help='solo el autotest sobre el fichero existente')
    ap.add_argument('--salida', default=DESTINO)
    a = ap.parse_args()
    if not a.autotest:
        if generar(a.salida):
            print('ERROR: fallos de pycel en inject_cache')
            return 1
    return autotest(a.salida)


if __name__ == '__main__':
    sys.exit(main())
