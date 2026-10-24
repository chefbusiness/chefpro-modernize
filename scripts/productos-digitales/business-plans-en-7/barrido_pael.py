#!/usr/bin/env python3
"""
barrido_pael.py — barrido INDEPENDIENTE de los 4 entregables del Paella Catering Business Plan Kit (EN) en
`astro-site/public/dl/paella-catering-business-plan/`: todas las celdas de texto, los literales de las fórmulas, los párrafos
y tablas del docx y los docProps, contra el vocabulario vetado RX_PAEL_NO y las cifras prohibidas (D128:
`mapas.cifras_prohibidas`), restos de español (`mapas.restos_espanol`, sin la categoría «oficio») y caracteres no
latinos; más recuentos y sha256 → `entregables_pael.json`. Copia de `../business-plans-en-6/barrido_bbq.py`. Corre en
el VPS. Sesión Claude Code, 24-oct-2026 (F2, paso 1; se usa en el paso 3).

    python3 bp7.py barrido
"""
import hashlib
import json
import os
import re
import sys
from collections import OrderedDict

import docx
import openpyxl

AQUI = os.path.dirname(os.path.abspath(__file__))
REPO = os.path.abspath(os.path.join(AQUI, '..', '..', '..'))
sys.path.insert(0, os.path.join(os.path.dirname(AQUI), 'business-plans-en'))
os.environ.setdefault('BP_CONJUNTO', 'pael')
import mapas as M                                                                  # noqa: E402

DIR = os.path.join(REPO, 'astro-site', 'public', 'dl', 'paella-catering-business-plan')
RX_LIT = re.compile(r'"((?:[^"]|"")*)"')


def piezas_xlsx(p):
    wb = openpyxl.load_workbook(p)
    out, n_f = [], 0
    for ws in wb.worksheets:
        out.append(('%s (pestaña)' % ws.title, ws.title))
        for row in ws.iter_rows():
            for c in row:
                v = c.value
                if not isinstance(v, str):
                    continue
                if c.data_type == 'f':
                    n_f += 1
                    for lit in RX_LIT.findall(v):
                        if any(ch.isalpha() for ch in lit):
                            out.append(('%s!%s (fórmula)' % (ws.title, c.coordinate), lit))
                else:
                    out.append(('%s!%s' % (ws.title, c.coordinate), v))
        for dv in ws.data_validations.dataValidation:
            for t in (dv.error, dv.errorTitle, dv.prompt, dv.promptTitle):
                if t:
                    out.append(('%s DV %s' % (ws.title, dv.sqref), t))
    pr = wb.properties
    for campo in ('title', 'subject', 'keywords', 'description', 'category', 'creator'):
        if getattr(pr, campo):
            out.append(('docProps %s' % campo, getattr(pr, campo)))
    return out, OrderedDict([('hojas', len(wb.worksheets)), ('formulas', n_f)])


def piezas_docx(p):
    d = docx.Document(p)
    out = [('párrafo %d' % i, x.text) for i, x in enumerate(d.paragraphs) if x.text.strip()]
    for t, tab in enumerate(d.tables):
        for r, row in enumerate(tab.rows):
            for c, cel in enumerate(row.cells):
                if cel.text.strip():
                    out.append(('tabla %d [%d,%d]' % (t, r, c), cel.text))
    cp = d.core_properties
    for campo in ('title', 'subject', 'keywords', 'comments', 'category', 'author'):
        if getattr(cp, campo):
            out.append(('core %s' % campo, getattr(cp, campo)))
    palabras = sum(len(x.text.split()) for x in d.paragraphs)
    return out, OrderedDict([('parrafos', len(d.paragraphs)), ('con_texto', sum(1 for x in d.paragraphs if x.text.strip())),
                             ('heading1', sum(1 for x in d.paragraphs if x.style.name == 'Heading 1')),
                             ('palabras', palabras)])


def main():
    res, fallos, total = OrderedDict(), [], 0
    for f in sorted(os.listdir(DIR)):
        p = os.path.join(DIR, f)
        if f.endswith('.xlsx'):
            piezas, info = piezas_xlsx(p)
        elif f.endswith('.docx'):
            piezas, info = piezas_docx(p)
        else:
            continue
        mal = OrderedDict([('prohibido', []), ('espanol', []), ('no_latinos', [])])
        for donde, t in piezas:
            for motivo, frag in M.cifras_prohibidas(t):
                mal['prohibido'].append('%s: %s «%s»' % (donde, motivo, frag))
            r = [x for x in M.restos_espanol(t) if not x[0].startswith('oficio')]
            if r:
                mal['espanol'].append('%s: %s' % (donde, r[0][0]))
            if M.no_latinos(t):
                mal['no_latinos'].append(donde)
        total += len(piezas)
        info['piezas'] = len(piezas)
        info['sha256'] = hashlib.sha256(open(p, 'rb').read()).hexdigest()
        info.update((k, len(v)) for k, v in mal.items())
        res[f] = info
        for k, v in mal.items():
            fallos += ['%s · %s' % (f, x) for x in v]
    out = OrderedDict([('carpeta', os.path.relpath(DIR, REPO)), ('piezas', total), ('fallos', len(fallos)),
                       ('ficheros', res)])
    with open(os.path.join(AQUI, 'entregables_pael.json'), 'w', encoding='utf-8') as fh:
        json.dump(out, fh, ensure_ascii=False, indent=1)
        fh.write('\n')
    for x in fallos[:40]:
        print('ERROR', x)
    for f, i in res.items():
        print('%-40s %s' % (f, ' · '.join('%s %s' % (k, (v[:16] if k == 'sha256' else v)) for k, v in i.items())))
    print('barrido: %d ficheros · %d piezas · %d fallos (RX_PAEL_NO / cifras prohibidas / español / no latinos)'
          % (len(res), total, len(fallos)))
    return 1 if fallos else 0


if __name__ == '__main__':
    sys.exit(main())
