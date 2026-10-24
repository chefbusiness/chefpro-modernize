#!/usr/bin/env python3
"""
fuente_pael.py — material de CONTENIDO para los redactores: el checklist ES del Paellero 2.2 (83 tareas, desde
`../paellero-2_2/datos_checklist.py`, sin abrir el xlsx), los párrafos ES 2.2 del plan en las 61 ranuras del docx EN
(`../paellero-2_2/textos/pae_es_{a,b}.json`, con sus {TOKENS}), y las tablas y textos del manual técnico (escalado de
paelleras, cubrir 50/100/200 invitados, 12 paellas, cronograma, alérgenos), la carta de 12 paellas y arroces, las 10
experiencias y el contrato (python-docx) en `fuente_es_pael.json`. El circuito de texto corre sobre el FOOD TRUCK ES (las
ranuras); el contenido de paella sale de aquí y de research (D119, D125, D126). Sus CIFRAS no se usan (D128). Corre en el
VPS (python-docx). Copia de `../business-plans-en-6/fuente_bbq.py`. Sesión Claude Code, 24-oct-2026 (F2, paso 1).

    python3 bp7.py fuente
"""
import datetime
import hashlib
import importlib.util
import json
import os
from collections import OrderedDict

import docx

AQUI = os.path.dirname(os.path.abspath(__file__))
REPO = os.path.abspath(os.path.join(AQUI, '..', '..', '..'))
PAE = os.path.join(os.path.dirname(AQUI), 'paellero-2_2')
DIR = os.path.join(REPO, 'astro-site', 'public', 'dl', 'plan-negocio-paellero-eventos')
MANUAL = 'manual-tecnico-paellero-profesional.docx'
CARTA = 'carta-12-paellas-arroces.docx'
EXPERIENCIAS = '10-experiencias-tematicas-paellero.docx'
CONTRATO = 'modelos-contrato-paellero-eventos.docx'
EQUIPO = 'catalogo-equipamiento-utensilios-menaje-paellero.docx'
# F1-inventario §2: tareas solo España (fase, nº de tarea 1-based) — fuera del EN; y las parciales (lo español, fuera)
SOLO_ES = {(1, 3), (1, 4), (1, 5), (1, 8), (1, 10), (1, 11), (2, 3), (4, 2), (4, 7), (5, 14)}
PARCIALES = {(1, 7): 'gestoría y Veri*factu → contador', (1, 9): 'modelos 202 / 130 / 111 → estimated taxes',
             (1, 12): 'LSSI → términos de la web', (2, 10): 'regla de los 15 kg → fire code local',
             (2, 11): 'empresa instaladora (ITC-ICG) → fire marshal', (2, 14): 'época de peligro por CCAA → burn ban',
             (2, 16): 'hojas de reclamaciones → fuera', (5, 10): 'Hosply → ezCater'}
FASE_EN = OrderedDict([
    (1, 'P1 Business Setup (1-2 → LLC; 6 banco; 7 y 9 → contador y estimated taxes; 12 → términos de la web) · P1 contrato '
        '(F6·4) y plan HACCP (F2·4)'),
    (2, 'P2 Propane & Permits (1 commissary; 2 → county catering / health permit + limited food preparation; 4 → plan '
        'HACCP con arroz TCS; 7-9 → GL sin exclusión de llama abierta, commercial auto, inland marine; 10-13 → propano, '
        'cilindros en la van, extintores; 14 → burn ban y leña; 15 → fire marshal y venue; 16 → residuos y aceite)'),
    (3, 'P3 Equipment (1-4 → paelleras, quemadores, leña con permiso, cilindros; 5-10 → utensilios, coolers y sondas, '
        'kit de la estación, EPI; 11-13 → arroz bomba, marisco con trazabilidad y shellstock tags, azafrán y despensa; 14 '
        '→ van y software)'),
    (4, 'P4 Staff (1 → W-2; 3-5 → food handler cards, alérgenos, OSHA quemaduras / propano / cargas; 8 → segundo paella '
        'chef; 9 → guion del show (socarrat reveal); 10 → fichas por tamaño de paellera)'),
    (5, 'P5 Marketing (1-7; 8 WhatsApp → texto / SMS; 9 → The Knot / WeddingWire / Zola; 10 → ezCater; 11-13 → planners, '
        'empresas, venues; 15-16 → dossiers, referidos, tasting de lanzamiento)'),
    (6, 'P6 First 90 Days (1-6 briefing, visita, depósito, contrato, alérgenos, plan B de viento y lluvia; 7-8 → pedidos '
        'y la víspera en la commissary; 9-12 → carga, llegada 2-3 h antes, cocción y servicio, recogida; 13-15 → cobro, '
        'reseña, costes reales frente a la calculadora)'),
])
USO = ('CONTENIDO de paella catering para los redactores (D119, D125, D126). NO se traduce tal cual: se condensa y se '
       'adapta a EE. UU. con research (F1-research-us-delta.md §1, §5, §6, §8 y el glosario §9). Las CIFRAS del ES NO se '
       'usan (D128: 157.920, 105 eventos/año, 94 invitados, 16 €/invitado, 22,4 %…): las del caso salen del libro EN '
       '(tokens). Fuera: tareas marcadas SOLO_ES y lo español de las PARCIALES, «España», «ChefBusiness», normativa '
       'española (RD, Reglamento (CE), ITC-ICG, ADR, AESAN, APPCC, CCAA), proveedores y marcas españolas (Garcima, '
       'Vaello, GGM, Arcos), importes en €, «autónomo», «fijo discontinuo» y «freelance» (en US: part-time hourly W-2 '
       'staff), fiestas populares y contrato menor, el anexo CCAA y la deuda del ES («Paella valenciana DO»: la DO '
       'ampara el arroz, no el plato). Medidas US: pulgadas para las paelleras con cm entre paréntesis, oz de arroz '
       'crudo, galones y cuartos de caldo, °F con el mínimo del Food Code nombrado como tal; azafrán en gramos')


def sha(p):
    return hashlib.sha256(open(p, 'rb').read()).hexdigest()


def checklist():
    p = os.path.join(PAE, 'datos_checklist.py')
    spec = importlib.util.spec_from_file_location('datos_checklist_pae', p)
    m = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(m)
    fases = OrderedDict()
    for k, (hoja, titulo, tareas) in enumerate(m.FASES, 1):
        lista = []
        for j, (t, resp, plazo, nota) in enumerate(tareas, 1):
            x = OrderedDict([('n', j), ('tarea', t), ('responsable', resp), ('plazo', plazo), ('nota', nota)])
            if (k, j) in SOLO_ES:
                x['marca'] = 'SOLO_ES (fuera)'
            elif (k, j) in PARCIALES:
                x['marca'] = 'PARCIAL: fuera lo español (%s)' % PARCIALES[(k, j)]
            lista.append(x)
        fases[hoja] = OrderedDict([('titulo', titulo), ('fase_en_sugerida', FASE_EN[k]), ('tareas', lista)])
    return p, m.TITULO, fases


def ranuras():
    out, fuentes = OrderedDict(), []
    for f in ('a', 'b'):
        p = os.path.join(PAE, 'textos', 'pae_es_%s.json' % f)
        fuentes.append(p)
        out.update(json.load(open(p, encoding='utf-8')))
    return fuentes, OrderedDict(sorted(out.items(), key=lambda kv: int(kv[0])))


def parrafos(nombre, minimo=3):
    p = os.path.join(DIR, nombre)
    d = docx.Document(p)
    pars = [OrderedDict([('i', i), ('estilo', x.style.name), ('texto', x.text)]) for i, x in enumerate(d.paragraphs)
            if len(x.text.strip()) >= minimo and not set(x.text.strip()) <= set('━─—-· ')]
    tablas = [[[c.text for c in row.cells] for row in t.rows] for t in d.tables]
    return p, OrderedDict([('parrafos', pars), ('tablas', tablas)])


def main():
    pc, titulo, fases = checklist()
    pr, docx_es = ranuras()
    docs = OrderedDict()
    fuentes = [pc] + pr
    for clave, nombre in (('manual_tecnico', MANUAL), ('carta_paellas', CARTA), ('experiencias_10', EXPERIENCIAS),
                          ('contrato', CONTRATO), ('catalogo_equipo', EQUIPO)):
        p, contenido = parrafos(nombre)
        docs[clave] = contenido
        fuentes.append(p)
    out = OrderedDict([
        ('meta', OrderedDict([
            ('generado', datetime.datetime.now(datetime.timezone.utc).isoformat(timespec='seconds')),
            ('script', 'scripts/productos-digitales/business-plans-en-7/fuente_pael.py'),
            ('uso', USO),
            ('fuentes', OrderedDict((os.path.relpath(x, REPO), sha(x)) for x in fuentes)),
        ])),
        ('checklist_es', OrderedDict([('titulo', titulo), ('fases', fases)])),
        ('docx_es_por_ranura', docx_es),
    ] + list(docs.items()))
    dest = os.path.join(AQUI, 'fuente_es_pael.json')
    with open(dest, 'w', encoding='utf-8') as fh:
        json.dump(out, fh, ensure_ascii=False, indent=1)
        fh.write('\n')
    n = sum(len(f['tareas']) for f in fases.values())
    fuera = sum(1 for f in fases.values() for t in f['tareas'] if t.get('marca', '').startswith('SOLO_ES'))
    print('fuente_es_pael.json: %d fases, %d tareas (%d fuera, %d parciales) · docx %d ranuras · %s' % (
        len(fases), n, fuera, len(PARCIALES), len(docx_es),
        ' · '.join('%s %d párrafos / %d tablas' % (k, len(v['parrafos']), len(v['tablas'])) for k, v in docs.items())))
    return 0


if __name__ == '__main__':
    raise SystemExit(main())
