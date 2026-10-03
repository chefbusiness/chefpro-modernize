#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
gates_en.py — Gates F2 del Restaurant Staff Scheduling Kit Pro (EN), SPEC §8 (G1-G8), sobre la salida de aplicar_en.py.
Copia adaptada de `haccp-kit/gates_en.py`. Solo LEE: los 9 xlsx EN, los 9 xlsx ES publicados (dl/kit-gestion-personal,
solo lectura), censo_es.json y mapas.py. Ningún recuento está escrito a mano: todo sale de censo_es.json, de los xlsx
ES o de mapas.py; las únicas cifras propias son las que la SPEC fija (G6: D8, D10, D20).

    python3 gates_en.py                               # contra astro-site/public/dl/restaurant-schedule-templates/
    python3 gates_en.py --dir /tmp/staff-dry          # contra un dry-run
    python3 gates_en.py --solo G1,G2 [--mes October] [--json informe.json]
    python3 gates_en.py --autotest [--dir …]          # inyecta defectos en una COPIA y exige que se detecten

Gates (SPEC §8):
  G1 paridad: 30 hojas por §2.2; fórmula a fórmula = ES tras HOJAS + CLAVES salvo los 10 patrones de FORMULAS_EN;
     fórmulas con hoja; DV (mismo sqref y nº de celdas, flags, tipo; listas por CLAVES salvo D11) + las 4 DV_NUEVAS
     exactas; CF (sqref, tipo, prioridad, relleno, fórmula y `text`); merges; áreas y títulos de impresión; paneles;
     columnas ocultas; protección (21 hojas, sin contraseña, mismos flags); verdes y desbloqueadas (+3 de D10);
     0 gráficos e imágenes;
  G2 restos: cero español, «€», caracteres fuera de la lista blanca, coma decimal, horas de 24 h, normativa ES
     (RX_NORMA) e IDs internos (también los «kitgp-v2 · …» de los títulos de DV) en valores, literales, DV (ítem a
     ítem), CF, formatos, pies y docProps; «°C» solo entre paréntesis tras su °F; citas a otro fichero de UN nivel
     (file NN "nombre", sin «((» ni «)))»); DV con título ≤ 32 y mensaje ≤ 255 (límites de Excel); DV_MENSAJES_EN;
  G3 formato: paperSize 1 en las 30; fechas del censo con numFmtId 14 o «m/d»; horas con 18; «Shifts» C5:D12 con
     «h:mm AM/PM;;»; ningún otro formato de fecha; pie D18 en las 30; línea de marca D18 donde el ES la tenía;
     versión D18; docProps; Instructions!B2 = título interior D5; AJUSTE_TEXTO con ajuste y alto de fila;
  G4 claves: cada ítem de DV ∈ CLAVES (o D11); cada celda escrita en un rango con DV de lista es un ítem; las tablas
     de tipos de negocio (03, B02) = la lista de su DV; cada número de una DV numérica, dentro de sus límites;
  G5 CF: tokens EN inyectivos por rango, todos ∈ CLAVES (o sin letras) y mapas.cruzar_censo (ningún mensaje cambia
     de color);
  G6 cálculo: <v> en toda <f>, cero errores; caché EN = caché ES (veredictos por CLAVES, números iguales) salvo
     las cifras que la SPEC cambia (D20), que se comprueban con su valor: Shifts!E5:E12 = 8,8,8,9,16,0,0,0; B02 = 7
     FTE, 72,800, 23,356.67, 32.1 %, EXCELLENT; 03 Staffing Forecast!B22 = 23,356.67; la frase de B02
     Instructions!B16; casos FLSA de D10 en copias con la semana en lunes (B3 = 2): 6 × 9 h → 14 h con F3 vacío
     y con F3 = 8; 4 × 10 h → 0 y 8; 6 × 9 h de miércoles a lunes → 5 h (la semana corta en lunes);
  G7 reglas: VALORES_EN, PARAMETROS, DV_NUEVAS y FORMULAS_EN = SPEC §3 (mapas.REGLAS_US);
  G8 censo-entregables.py --fail (Letter) y gate-no-latinos.py en 0.
Sale con código 1 si algún gate falla.
"""
import argparse
import datetime
import html
import json
import logging
import os
import re
import shutil
import subprocess
import sys
import tempfile
import unicodedata
import zipfile
from collections import OrderedDict, Counter

AQUI = os.path.dirname(os.path.abspath(__file__))
SCRIPTS = os.path.dirname(AQUI)
REPO = os.path.dirname(os.path.dirname(SCRIPTS))
sys.path.insert(0, AQUI)

import openpyxl                                          # noqa: E402

import mapas                                             # noqa: E402
import aplicar_en as A                                   # noqa: E402

logging.disable(logging.CRITICAL)

ORIGEN = A.ORIGEN
CENSO_SCRIPT = os.path.join(SCRIPTS, 'censo-entregables.py')
NO_LATINOS = os.path.join(SCRIPTS, 'gate-no-latinos.py')
TOL = 0.006
ERRORES = ('#REF!', '#VALUE!', '#NAME?', '#DIV/0!', '#N/A', '#NUM!', '#NULL!')
CLAVES = mapas.CLAVES
CLAVES_EN = set(CLAVES.values())
D11_ITEMS = {x for _es, en in mapas.DV_LISTAS_ESPECIALES.values() for x in en.strip('"').split(',')}

SIGNOS_ES = set('¿¡«»€ºª')
RX_ID_INTERNO = re.compile(r'\[(?:fuente|estimado|derivado|criterio)[^\]]*\]|\bSPEC\b|\((?:D|I)\d{1,2}\)|'
                           r'\bD\d{1,2}\b(?=[:,)])|\bTEC-\d+\b|'         # TEC-nn: ticket interno del ES (B01)
                           r'\b(?:kitgp|grupo[A-Z])-v\d+[a-z]?\b')           # IDs de las DV del ES (revisión final)
# Citas a otro fichero: file NN "nombre", UN nivel (revisión final: «file 03 (Restaurant … (Payroll & Labor %))»)
RX_CITA_ANIDADA = re.compile(r'\(\(|\)\)\)|\bfile (?:0\d|BONUS-0\d) \(')
LIMITE_DV = {'promptTitle': 32, 'errorTitle': 32, 'prompt': 255, 'error': 255}     # Excel
RX_TILDE = re.compile(r'[áéíóúüñÁÉÍÓÚÜÑ]')
LISTA_BLANCA_TILDE = {'café', 'cafés', 'entrée', 'entrées', 'sauté', 'sautéed', 'purée', 'jalapeño', 'jalapeños',
                      'crème', 'brûlée', 'résumé', 'résumés', 'naïve'}
FUNCIONALES_ES = {
    'de', 'del', 'las', 'el', 'los', 'y', 'que', 'con', 'para', 'por', 'una', 'unos', 'unas', 'es', 'al', 'se',
    'su', 'sus', 'como', 'pero', 'muy', 'cuando', 'donde', 'porque', 'este', 'esta', 'esto', 'estos', 'desde',
    'hasta', 'sobre', 'entre', 'tu', 'tus', 'te', 'mis', 'hoja', 'pestaña', 'celda', 'celdas', 'sin', 'ejemplo',
    'cada', 'debe', 'hay', 'puede', 'si', 'lo', 'le', 'les', 'todo', 'toda', 'todos', 'nunca', 'siempre',
}
OFICIO_ES = {
    'turno', 'turnos', 'cuadrante', 'jornada', 'jornadas', 'horas', 'hora', 'empleado', 'empleados', 'empleada',
    'vacaciones', 'nomina', 'nominas', 'plantilla', 'plantillas', 'semana', 'semanas', 'lunes', 'martes',
    'miercoles', 'jueves', 'viernes', 'sabado', 'domingo', 'descanso', 'camarero', 'camarera', 'cocinero',
    'cocinera', 'ayudante', 'encargado', 'encargada', 'barra', 'dia', 'dias', 'fecha', 'fechas', 'alta', 'baja',
    'contrato', 'salario', 'bruto', 'coste', 'cotizacion', 'pagas', 'festivo', 'permiso', 'solicitud',
    'solicitudes', 'saldo', 'cobertura', 'evaluacion', 'ficha', 'desempeno', 'puesto', 'puestos', 'vencimiento',
    'vencimientos', 'caducidad', 'carnet', 'formacion', 'cubiertos', 'ventas', 'negocio', 'cocina', 'hosteleria',
    'calculadora', 'servicio', 'servicios', 'temporada', 'ausencia', 'ausencias', 'menor', 'menores', 'hecho',
    'plazo', 'notas', 'responsable', 'tarea', 'tareas', 'categoria', 'empresa', 'gestoria', 'convenio',
    'trabajador', 'trabajadores', 'nombre', 'apellidos', 'telefono', 'correo', 'aviso', 'alerta', 'semaforo',
    'objetivo', 'importe', 'cambio', 'caja', 'arqueo', 'efectivo', 'fondo', 'temperatura', 'equipo', 'conforme',
    'incidencias', 'reservas', 'comida', 'cena', 'desayuno', 'manana', 'tarde', 'noche', 'partido', 'doble',
    'libre', 'firma', 'observaciones', 'pendiente', 'aprobado', 'denegado', 'resumen', 'registro', 'instrucciones',
    'mensual', 'semanal', 'anual', 'previsto', 'previstos', 'personal', 'sala',
}
OFICIO_ES -= {'personal'}                               # «personal» es también inglés («personal data»)
RX_URL = re.compile(r'https?://\S+|\b[\w-]+(?:\.[\w-]+)*\.(?:es|com|pro|gov|org|uk|co|example|net)\b[/\w.#-]*', re.I)
RX_PALABRA = re.compile(r"[A-Za-zÀ-ÖØ-öø-ÿ]+")
RX_STR = A.RX_STR
RX_NORMA = re.compile(A.extraer_textos.RX_NORMA.pattern)
RX_SIGLAS_ES = re.compile(r'\b(?:TPV|RLT|IVA|ERTE|CIF|NIF|SMI)\b')
RX_HORA_24 = re.compile(r'\b(?:0\d|1[3-9]|2[0-3]):[0-5]\d\b')
RX_COMA_DECIMAL = re.compile(r'(?<![\d,])\d+,\d{1,2}(?![\d])')


def _simbolos_es():
    """Lista blanca de caracteres no ASCII que no son letras: los del ES publicado (▸ → ✓ ✗ ⚠ ⛔ emojis…) +
    la tipografía del kit EN (§ ≤ ≥ − ± × ’ “ ”). Los signos españoles (¿ ¡ « » € º ª) nunca entran."""
    ok = set('▸→☐✓✗⚠·—–…°§≤≥−±×½⛔✅❌⏱↑↓⭐') | {'️'}
    ok.add('\U0001F4B5')              # 💵: el 💶 del ES (B01 «CAJA») pasa al billete de dólar en la edición US
    te = json.load(open(os.path.join(AQUI, 'textos_es.json'), encoding='utf-8'))
    textos = [c['es'] for c in te['cadenas']] + list(te['meta']['excluidas']['cadenas'])
    censo = json.load(open(os.path.join(AQUI, 'censo_es.json'), encoding='utf-8'))
    for info in censo['libros'].values():
        for hd in info['hojas_detalle'].values():
            textos += list(hd['literales']) + list(hd['literales_cf'])
    for t in textos:
        for ch in t:
            if ord(ch) > 127 and not unicodedata.category(ch).startswith('L') and ch not in SIGNOS_ES:
                ok.add(ch)
    return ok


SIMBOLOS_OK = _simbolos_es()


def no_permitidos(t):
    return sorted({ch for ch in t if ord(ch) > 127 and ch not in SIMBOLOS_OK
                   and not ('À' <= ch <= 'ÿ' and ch not in '×÷')})


def restos_espanol(t):
    out = []
    limpio = RX_URL.sub(' ', t)
    for ch in limpio:
        if ch in SIGNOS_ES:
            out.append(('signo «%s»' % ch, t[:80]))
    for m in RX_PALABRA.finditer(limpio):
        w = m.group(0)
        wl = w.lower()
        ctx = limpio[max(0, m.start() - 30):m.end() + 30].replace('\n', ' ')
        sin = (wl.replace('á', 'a').replace('é', 'e').replace('í', 'i').replace('ó', 'o').replace('ú', 'u')
               .replace('ñ', 'n'))
        if RX_TILDE.search(w) and wl not in LISTA_BLANCA_TILDE:
            out.append(('tilde/eñe «%s»' % w, ctx))
        elif wl in FUNCIONALES_ES and w not in ('Y', 'T', 'N', 'NO', 'SE', 'TE', 'SI', 'S'):   # Y/N/S son claves EN
            out.append(('funcional ES «%s»' % w, ctx))
        elif sin in OFICIO_ES and wl not in LISTA_BLANCA_TILDE:
            out.append(('oficio ES «%s»' % w, ctx))
    m = RX_COMA_DECIMAL.search(t)
    if m:
        out.append(('coma decimal «%s»' % m.group(0), t[:80]))
    m = RX_HORA_24.search(t)
    if m and not re.search(re.escape(m.group(0)) + r'\s*(?:AM|PM|a\.m\.|p\.m\.)', t):
        out.append(('hora de 24 h «%s»' % m.group(0), t[:80]))
    m = RX_SIGLAS_ES.search(t)
    if m:
        out.append(('sigla ES «%s»' % m.group(0), t[:80]))
    m = RX_NORMA.search(t)
    if m:
        out.append(('normativa ES «%s»' % m.group(0), t[:80]))
    return out


def celsius_sueltos(t):
    """D16: todo «°C» de texto va entre paréntesis detrás de su cifra en °F («41 °F (5 °C)»)."""
    malos = []
    for m in re.finditer('°C', t):
        antes = t[:m.start()]
        abre = antes.rfind('(')
        if abre < 0 or ')' in antes[abre:]:
            malos.append(t[max(0, m.start() - 25):m.end() + 5])
            continue
        if '°F' not in antes[abre:] and '°F' not in antes[max(0, abre - 40):abre]:
            malos.append(t[max(0, abre - 25):m.end() + 5])
    return malos


# ==========================================================================
# Utilidades
# ==========================================================================
class Resultado:
    def __init__(self):
        self.gates = OrderedDict()
        self.actual = None

    def gate(self, nombre):
        self.actual = nombre
        self.gates.setdefault(nombre, {'fallos': [], 'datos': OrderedDict()})
        return self

    def fallo(self, msg):
        self.gates[self.actual]['fallos'].append(msg)

    def dato(self, k, v):
        self.gates[self.actual]['datos'][k] = v


_CACHE = {}


def cargar(path, data_only=False):
    k = (path, data_only, os.path.getmtime(path))
    if k not in _CACHE:
        _CACHE[k] = openpyxl.load_workbook(path, data_only=data_only)
    return _CACHE[k]


def libros(carpeta):
    for es in mapas.LIBROS_ES:
        en = mapas.FICHEROS[es]
        yield es, en, A.extraer_textos.corto(es), os.path.join(ORIGEN, es), os.path.join(carpeta, en)


def fichero_en(carpeta, corto):
    return [pen for es, en, c, pes, pen in libros(carpeta) if c == corto][0]


def lits(f):
    return [x[1:-1].replace('""', '"') for x in RX_STR.findall(f or '')]


def celdas(sqref):
    out = []
    for rg in str(sqref).split():
        out.extend(A.rango(rg))
    return out


def es_verde(c):
    return (c.fill is not None and c.fill.fgColor is not None and isinstance(c.fill.fgColor.rgb, str)
            and c.fill.fgColor.rgb.upper().endswith(A.VERDE))


def celdas_reales(ws):
    for row in ws.iter_rows():
        for c in row:
            if c.__class__.__name__ != 'MergedCell':
                yield c


def _pagina(ws):
    ps = ws.page_setup
    fit = ws.sheet_properties.pageSetUpPr.fitToPage if ws.sheet_properties.pageSetUpPr else None
    return (ps.orientation, ps.fitToWidth, ps.fitToHeight, fit, str(ws.print_title_rows), str(ws.print_title_cols))


def _area(ws):
    pa = ws.print_area
    return pa.split('!')[-1] if pa else None


def _proteccion(ws):
    p = ws.protection
    return (bool(p.sheet), bool(p.password), p.formatColumns, p.formatRows, p.insertRows, p.deleteRows, p.sort,
            p.autoFilter, p.selectLockedCells, p.selectUnlockedCells)


def _cf_firma(cf, regla, trans, trans_text):
    dxf = regla.dxf
    color = None
    if dxf is not None:
        color = (str(dxf.fill.fgColor.rgb) if dxf.fill is not None and dxf.fill.fgColor is not None else None,
                 str(dxf.font.color.rgb) if dxf.font is not None and dxf.font.color is not None else None)
    return (str(cf.sqref), regla.type, regla.operator, regla.priority, bool(regla.stopIfTrue),
            tuple(trans(x) for x in (regla.formula or [])), trans_text(regla.text) if regla.text else regla.text, color)


def dv_nuevas_de(corto, h_es):
    return OrderedDict((sq, v) for (c, h, sq), v in mapas.DV_NUEVAS.items() if c == corto and h == h_es)


def dv_en_esperada(corto, h_es, de, T):
    """(type, operator, formula1, formula2) EN de una DV ES (misma lógica que aplicar_en, recalculada aquí)."""
    sq, f1 = str(de.sqref), de.formula1 or ''
    esp = mapas.DV_LISTAS_ESPECIALES.get((corto, h_es, sq))
    if esp:
        return de.type, de.operator, esp[1], de.formula2
    if de.type == 'list' and f1.startswith('"'):
        items = []
        for x in f1.strip('"').split(','):
            items.append(CLAVES[x] if mapas.necesita_clave(x) else x)
        return de.type, de.operator, mapas.dv_lista(items), de.formula2
    if de.type == 'custom' and f1:
        return de.type, de.operator, T.formula(f1), de.formula2
    return (de.type, de.operator, T.refs(f1) if f1 else de.formula1,
            T.refs(de.formula2) if de.formula2 else de.formula2)


def serial(v):
    """Valor de celda → número (fechas y horas a serie de Excel) o None si no es numérico."""
    if isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        return v
    if isinstance(v, datetime.datetime):
        return (v - datetime.datetime(1899, 12, 30)).total_seconds() / 86400.0
    if isinstance(v, datetime.time):
        return (v.hour * 3600 + v.minute * 60 + v.second) / 86400.0
    if isinstance(v, datetime.date):
        return (v - datetime.date(1899, 12, 30)).days
    return None


# ==========================================================================
# G1 — paridad estructural
# ==========================================================================
def gateG1(R, carpeta, datos):
    R.gate('G1 paridad')
    censo = datos['censo']
    tot = Counter()
    nuevas_verdes = {(c, mapas.HOJAS[h], coord) for (c, h, coord) in mapas.CELDAS_VERDES_NUEVAS}
    for es, en, corto, pes, pen in libros(carpeta):
        if not os.path.isfile(pen):
            R.fallo('%s: no existe' % en)
            continue
        ce = censo['libros'][es]
        wes, wen = cargar(pes), cargar(pen)
        esperado = [mapas.HOJAS[h] for h in ce['hojas']]
        if wen.sheetnames != esperado or wes.sheetnames != ce['hojas']:
            R.fallo('%s: hojas %r ≠ %r' % (en, wen.sheetnames, esperado))
            continue
        tot['hojas'] += len(esperado)
        mapa = OrderedDict((h, mapas.HOJAS[h]) for h in wes.sheetnames)
        T = A.Transformador(mapa)
        for h_es in wes.sheetnames:
            ws_es, ws_en = wes[h_es], wen[mapa[h_es]]
            hd = ce['hojas_detalle'][h_es]
            lugar = '%s:%s' % (en, ws_en.title)
            f_es = {c.coordinate: c.value for c in ws_es._cells.values() if c.data_type == 'f'}
            f_en = {c.coordinate: c.value for c in ws_en._cells.values() if c.data_type == 'f'}
            if len(f_es) != hd['celdas_formula']:
                R.fallo('%s: el ES ya no cuadra con el censo (%d ≠ %d fórmulas)' % (lugar, len(f_es), hd['celdas_formula']))
            if set(f_en) != set(f_es):
                R.fallo('%s: celdas con fórmula distintas: solo ES %s · solo EN %s'
                        % (lugar, sorted(set(f_es) - set(f_en))[:5], sorted(set(f_en) - set(f_es))[:5]))
            con_hoja = delta_hoja = 0
            for k, fe in f_es.items():
                fn = f_en.get(k)
                if fn is None:
                    continue
                tot['formulas'] += 1
                try:
                    esperado_f, es_pat = A.formula_en(fe, ws_es[k].row, T)
                except A.Aborta as e:
                    esperado_f, es_pat = '<<%s>>' % e, False
                tot['formulas_patron_en'] += es_pat
                if es_pat:                     # D10: «Monthly OT Summary»!G ya no cita «Time Log» (coste = B × tarifa × recargo)
                    delta_hoja += bool(A.extraer_textos.refs_hoja(esperado_f)) - bool(A.extraer_textos.refs_hoja(fe))
                if esperado_f != fn:
                    R.fallo('%s!%s: fórmula EN %r ≠ esperada %r' % (lugar, k, fn[:80], esperado_f[:80]))
                if A.extraer_textos.refs_hoja(fn):
                    con_hoja += 1
                    for h in A.extraer_textos.refs_hoja(fn):
                        if h not in wen.sheetnames:
                            R.fallo('%s!%s cita la hoja inexistente %r' % (lugar, k, h))
            tot['formulas_con_hoja'] += con_hoja
            tot['formulas_con_hoja_delta_patrones'] += delta_hoja
            if con_hoja != hd['formulas_con_hoja'] + delta_hoja:
                R.fallo('%s: %d fórmulas con referencia a hoja ≠ censo %d %+d (patrones FORMULAS_EN)'
                        % (lugar, con_hoja, hd['formulas_con_hoja'], delta_hoja))
            # DV: las del ES (transformadas) + las DV_NUEVAS exactas
            cen_dv = {d['sqref']: d for d in hd['dv']}
            nuevas = dv_nuevas_de(corto, h_es)
            dves = sorted(ws_es.data_validations.dataValidation, key=lambda d: str(d.sqref))
            dven_all = ws_en.data_validations.dataValidation
            extra = [d for d in dven_all if str(d.sqref) in nuevas]
            dven = sorted([d for d in dven_all if d not in extra], key=lambda d: str(d.sqref))
            if len(dves) != len(dven) or len(dves) != len(hd['dv']) or len(extra) != len(nuevas):
                R.fallo('%s: DV ES %d, EN %d (+%d nuevas de %d), censo %d' % (lugar, len(dves), len(dven), len(extra),
                                                                            len(nuevas), len(hd['dv'])))
            for de, dn in zip(dves, dven):
                tot['dv'] += 1
                tot['dv_celdas'] += len(celdas(dn.sqref))
                firma = lambda d: (str(d.sqref), bool(d.allow_blank), bool(d.showDropDown),  # noqa: E731
                                   bool(d.showErrorMessage), bool(d.showInputMessage), d.errorStyle)
                if firma(de) != firma(dn):
                    R.fallo('%s: DV %s cambia (%r → %r)' % (lugar, de.sqref, firma(de), firma(dn)))
                cen = cen_dv.get(str(de.sqref))
                if cen is None or len(celdas(dn.sqref)) != cen['n_celdas']:
                    R.fallo('%s: DV %s fuera del censo o con otro nº de celdas' % (lugar, dn.sqref))
                    continue
                for attr in ('error', 'errorTitle', 'prompt', 'promptTitle'):
                    if bool(getattr(de, attr)) != bool(getattr(dn, attr)):
                        R.fallo('%s: DV %s pierde/gana %s' % (lugar, dn.sqref, attr))
                try:
                    esp = dv_en_esperada(corto, h_es, de, T)
                except (KeyError, A.Aborta) as e:
                    esp = ('<<%s>>' % e,)
                got = (dn.type, dn.operator, dn.formula1, dn.formula2)
                if got != esp:
                    R.fallo('%s: DV %s %r ≠ esperada %r' % (lugar, dn.sqref, got, esp))
                if (corto, h_es, str(de.sqref)) in mapas.DV_LISTAS_ESPECIALES:
                    tot['dv_d11'] += 1
            for dn in extra:
                tipo, op, f1, f2, msg = nuevas[str(dn.sqref)]
                tot['dv_nuevas'] += 1
                if (dn.type, dn.operator, dn.formula1, dn.formula2) != (tipo, op, f1, f2) or dn.prompt != msg \
                        or dn.error != msg or not dn.showErrorMessage:
                    R.fallo('%s: DV nueva %s = %r ≠ DV_NUEVAS' % (lugar, dn.sqref, (dn.type, dn.formula1, dn.formula2)))
            # CF
            cfe = sorted((_cf_firma(cf, r, T.formula, lambda t: A.clave_en(t, 'CF')) for cf in ws_es.conditional_formatting
                          for r in cf.rules), key=repr)
            cfn = sorted((_cf_firma(cf, r, lambda x: x, lambda t: t) for cf in ws_en.conditional_formatting
                          for r in cf.rules), key=repr)
            tot['cf'] += len(cfn)
            if len(cfn) != len(hd['cf']):
                R.fallo('%s: %d CF EN ≠ censo %d' % (lugar, len(cfn), len(hd['cf'])))
            if cfe != cfn:
                dif = [(a, b) for a, b in zip(cfe, cfn) if a != b] or [(cfe[:1], cfn[:1])]
                R.fallo('%s: formato condicional distinto (ES transformado → EN): %r' % (lugar, dif[0]))
            # merges, protección, paneles, impresión, columnas ocultas
            mg = sorted(map(str, ws_en.merged_cells.ranges))
            if mg != sorted(hd['merges']):
                R.fallo('%s: combinaciones distintas' % lugar)
            tot['merges'] += len(mg)
            if _proteccion(ws_en) != _proteccion(ws_es) or ws_en.protection.password \
                    or bool(ws_en.protection.sheet) != hd['proteccion']['hoja']:
                R.fallo('%s: protección distinta del ES %r ≠ %r' % (lugar, _proteccion(ws_en), _proteccion(ws_es)))
            tot['hojas_protegidas'] += 1 if ws_en.protection.sheet else 0
            if ws_en.freeze_panes != hd['paneles']:
                R.fallo('%s: paneles %s ≠ %s' % (lugar, ws_en.freeze_panes, hd['paneles']))
            tot['paneles'] += 1 if ws_en.freeze_panes else 0
            if _pagina(ws_es) != _pagina(ws_en):
                R.fallo('%s: ajuste o títulos de impresión %r ≠ %r' % (lugar, _pagina(ws_es), _pagina(ws_en)))
            tot['titulos_impresion'] += 1 if ws_en.print_title_rows else 0
            ae = hd['area_impresion'].split('!')[-1] if hd['area_impresion'] else None
            if _area(ws_en) != ae:
                R.fallo('%s: área de impresión %s ≠ %s' % (lugar, _area(ws_en), ae))
            tot['areas_impresion'] += 1 if ae else 0
            ocultas = sorted(k for k, d in ws_en.column_dimensions.items() if d.hidden)
            if ocultas != hd['columnas_ocultas']:
                R.fallo('%s: columnas ocultas %s ≠ %s' % (lugar, ocultas, hd['columnas_ocultas']))
            tot['columnas_ocultas'] += len(ocultas)
            # verdes y desbloqueadas (+B3, D3, F3 de D10)
            extra_c = {coord for (c, h, coord) in nuevas_verdes if c == corto and h == ws_en.title}
            ue = {c.coordinate for c in celdas_reales(ws_es) if c.protection.locked is False}
            un = {c.coordinate for c in celdas_reales(ws_en) if c.protection.locked is False}
            ve = {c.coordinate for c in celdas_reales(ws_es) if es_verde(c)}
            vn = {c.coordinate for c in celdas_reales(ws_en) if es_verde(c)}
            if un != ue | extra_c:
                R.fallo('%s: desbloqueadas distintas (solo ES %s, solo EN %s)' % (lugar, sorted(ue - un)[:5],
                                                                                   sorted(un - ue - extra_c)[:5]))
            if vn != ve | extra_c:
                R.fallo('%s: verdes distintas (solo ES %s, solo EN %s)' % (lugar, sorted(ve - vn)[:5], sorted(vn - ve - extra_c)[:5]))
            if vn - un:
                R.fallo('%s: celdas verdes bloqueadas %s' % (lugar, sorted(vn - un)[:5]))
            tot['verdes'] += len(vn)
            tot['desbloqueadas'] += len(un)
        with zipfile.ZipFile(pen) as z:
            n = z.namelist()
        g = [x for x in n if x.startswith('xl/charts/') or x.startswith('xl/media/') or x.startswith('xl/drawings/')]
        if g:
            R.fallo('%s: gráficos/imágenes %s' % (en, g[:3]))
    T0 = censo['totales']
    n_pat = sum(p['n'] for info in censo['libros'].values() for hd in info['hojas_detalle'].values()
                for p in hd['patrones_formula'] if p['patron'] in mapas.FORMULAS_EN)
    n_ver = len(mapas.CELDAS_VERDES_NUEVAS)
    esperados = OrderedDict([
        ('hojas', T0['hojas']), ('formulas', T0['celdas_formula']), ('formulas_patron_en', n_pat),
        ('formulas_con_hoja', T0['formulas_con_hoja'] + tot['formulas_con_hoja_delta_patrones']), ('dv', T0['dv']), ('dv_nuevas', len(mapas.DV_NUEVAS)),
        ('dv_d11', len(mapas.DV_LISTAS_ESPECIALES)), ('dv_celdas', T0['dv_celdas']), ('cf', T0['cf']),
        ('merges', T0['merges']), ('hojas_protegidas', T0['hojas_protegidas']),
        ('areas_impresion', T0['areas_impresion']), ('titulos_impresion', T0['titulos_impresion']),
        ('paneles', T0['paneles']), ('columnas_ocultas', T0['columnas_ocultas']),
        ('verdes', T0['verdes'] + n_ver), ('desbloqueadas', T0['desbloqueadas'] + n_ver),
    ])
    for k, v in esperados.items():
        if tot[k] != v:
            R.fallo('total %s = %d ≠ esperado %d (censo)' % (k, tot[k], v))
    for k, v in sorted(tot.items()):
        R.dato(k, v)


# ==========================================================================
# G2 — restos y caracteres
# ==========================================================================
def piezas_de_texto(pen):
    wb = cargar(pen)
    for ws in wb.worksheets:
        yield 'hoja', ws.title, ws.title
        for c in ws._cells.values():
            if isinstance(c.value, str) and c.data_type != 'f':
                yield 'valor', '%s!%s' % (ws.title, c.coordinate), c.value
            if c.data_type == 'f':
                for x in lits(c.value):
                    yield 'literal', '%s!%s' % (ws.title, c.coordinate), x
            if isinstance(c.number_format, str) and c.number_format != 'General':
                yield 'formato', '%s!%s' % (ws.title, c.coordinate), c.number_format
        for dv in ws.data_validations.dataValidation:
            for attr in ('error', 'errorTitle', 'prompt', 'promptTitle'):
                v = getattr(dv, attr)
                if v:
                    yield 'dv', '%s DV %s %s' % (ws.title, dv.sqref, attr), v
            f1 = dv.formula1 or ''
            if f1.startswith('"'):
                for x in f1.strip('"').split(','):
                    if x:
                        yield 'dv-item', '%s DV %s' % (ws.title, dv.sqref), x
            elif dv.type == 'custom':
                for x in lits(f1):
                    yield 'literal', '%s DV %s' % (ws.title, dv.sqref), x
        for cf in ws.conditional_formatting:
            for regla in cf.rules:
                for f in regla.formula or []:
                    for x in lits(f):
                        yield 'literal', '%s CF %s' % (ws.title, cf.sqref), x
                if regla.text:
                    yield 'literal', '%s CF %s text' % (ws.title, cf.sqref), regla.text
        for parte in ('oddHeader', 'oddFooter', 'evenHeader', 'evenFooter', 'firstHeader', 'firstFooter'):
            hf = getattr(ws, parte)
            for pos in ('left', 'center', 'right'):
                t = getattr(hf, pos).text
                if t:
                    yield 'cabecera-pie', '%s %s.%s' % (ws.title, parte, pos), t
    p = wb.properties
    for k in ('title', 'subject', 'creator', 'keywords', 'description', 'category', 'lastModifiedBy'):
        v = getattr(p, k)
        if v:
            yield 'docprops', k, v
    with zipfile.ZipFile(pen) as z:
        estilos = z.read('xl/styles.xml').decode('utf-8')
    for m in re.finditer(r'<numFmt [^>]*formatCode="([^"]*)"', estilos):
        yield 'styles.xml', 'numFmt', html.unescape(m.group(1))


def gateG2(R, carpeta, datos):
    R.gate('G2 restos')
    n = Counter()
    for es, en, corto, pes, pen in libros(carpeta):
        for tipo, lugar, t in piezas_de_texto(pen):
            n[tipo] += 1
            donde = '%s %s' % (en, lugar)
            if '€' in t or re.search(r'\bEUR\b|\beuros?\b', t):
                R.fallo('%s: moneda € en %s: %r' % (donde, tipo, t[:70]))
            malos = no_permitidos(t)
            if malos:
                R.fallo('%s: carácter fuera de la lista blanca %r en %s: %r' % (donde, ''.join(malos), tipo, t[:60]))
            if tipo in ('formato', 'styles.xml'):
                if re.search(r'dd/mm|yyyy-mm-dd|€', t, re.I) or (re.search(r'hh?:mm', t, re.I) and 'AM/PM' not in t.upper()):
                    R.fallo('%s: formato %r' % (donde, t))
                continue
            m = RX_ID_INTERNO.search(t)
            if m:
                R.fallo('%s: ID interno %r en %s: …%s…' % (donde, m.group(0), tipo, t[max(0, m.start() - 30):m.end() + 10]))
            m = RX_CITA_ANIDADA.search(t)
            if m:
                R.fallo('%s: cita anidada %r en %s: …%s…' % (donde, m.group(0), tipo, t[max(0, m.start() - 40):m.end() + 30]))
            if tipo == 'dv':
                attr = lugar.rsplit(' ', 1)[-1]
                if len(t) > LIMITE_DV[attr]:
                    R.fallo('%s: DV %s de %d caracteres > %d (límite de Excel)' % (donde, attr, len(t), LIMITE_DV[attr]))
                n['dv_' + attr] += 1
            for motivo, ctx in restos_espanol(t):
                R.fallo('%s: %s · …%s…' % (donde, motivo, ctx))
            for s in celsius_sueltos(t):
                R.fallo('%s: °C sin su °F delante · …%s…' % (donde, s))
            if '°C' in t:
                n['textos_con_celsius_ok'] += 1
    # DV_MENSAJES_EN: el texto fijado por aparición está en su DV
    for (corto, h_es, sq, attr), v in mapas.DV_MENSAJES_EN.items():
        ws = cargar(fichero_en(carpeta, corto))[mapas.HOJAS[h_es]]
        dvs = [d for d in ws.data_validations.dataValidation if str(d.sqref) == sq]
        n['dv_mensajes_fijados'] += 1
        if len(dvs) != 1 or getattr(dvs[0], attr) != v:
            R.fallo('DV_MENSAJES_EN %s %s DV %s %s = %r ≠ %r' % (corto, ws.title, sq, attr,
                                                                  getattr(dvs[0], attr) if dvs else None, v))
    for k, v in sorted(n.items()):
        R.dato(k, v)


# ==========================================================================
# G3 — formato
# ==========================================================================
def fmt_esperado(fmt_es):
    return A.FORMATOS.get(fmt_es, fmt_es)


def gateG3(R, carpeta, datos, mes):
    R.gate('G3 formato')
    censo = datos['censo']
    n = Counter()
    rx_version = re.compile(r'^Version %s · %s \d{4} · %s · info@aichef\.pro$'
                            % (re.escape(A.VERSION), mes, re.escape(A.DESCRIPTION)))
    for es, en, corto, pes, pen in libros(carpeta):
        ce = censo['libros'][es]
        wb = cargar(pen)
        wes = cargar(pes)
        marcas_es = sum(1 for ws in wes.worksheets for c in ws._cells.values() if c.value == mapas.MARCA_ES)
        marcas_en = sum(1 for ws in wb.worksheets for c in ws._cells.values() if c.value == mapas.MARCA_EN)
        if marcas_en != marcas_es or marcas_es == 0:
            R.fallo('%s: %d líneas de marca D18 ≠ %d del ES' % (en, marcas_en, marcas_es))
        n['marcas'] += marcas_en
        for h_es, hd in ce['hojas_detalle'].items():
            ws = wb[mapas.HOJAS[h_es]]
            n['hojas'] += 1
            if ws.page_setup.paperSize not in (1, '1'):
                R.fallo('%s:%s: paperSize %s ≠ 1 (Letter)' % (en, ws.title, ws.page_setup.paperSize))
            fechas = {}
            for f in hd['fechas'] + hd['horas']:
                fechas[f['celda']] = fmt_esperado(f['formato'])
            if corto == '01' and h_es == 'Turnos':
                for coord in A.rango(A.TURNOS_RANGO_HORAS):
                    fechas[coord] = mapas.FMT_TURNOS
            for coord, fmt in fechas.items():
                c = ws[coord]
                fid = A.NUMFMT.get(fmt, 'x')
                if fid == 'x':
                    R.fallo('%s:%s!%s: formato ES %r sin equivalente US' % (en, ws.title, coord, fmt))
                    continue
                n['fmt_%s' % (fid if fid is not None else fmt)] += 1
                if fid is not None and c._style.numFmtId != fid:
                    R.fallo('%s:%s!%s: fecha/hora con numFmtId %s ≠ %s' % (en, ws.title, coord, c._style.numFmtId, fid))
                if fid is None and c.number_format != fmt:
                    R.fallo('%s:%s!%s: formato %r ≠ %r' % (en, ws.title, coord, c.number_format, fmt))
            for c in ws._cells.values():
                f = c.number_format or ''
                es_fecha = c._style.numFmtId in (14, 18, 22) or f in (A.FMT_FECHA_CORTA, mapas.FMT_TURNOS) or \
                    re.search(r'[dmy]/[dmy]|h:mm', f)
                if es_fecha and c.coordinate not in fechas:
                    R.fallo('%s:%s!%s: formato de fecha/hora %r en una celda que no lo era' % (en, ws.title, c.coordinate, f))
            for k in hd['cabeceras_pies']:
                obj, pos = k.split('.')
                v = getattr(getattr(ws, obj), pos).text
                if v != A.PIE_EN:
                    R.fallo('%s:%s: %s = %r ≠ pie D18' % (en, ws.title, k, v))
                else:
                    n['pies'] += 1
        ins = wb[A.HOJA_INSTR_EN]
        vers = [c.value for c in ins._cells.values() if isinstance(c.value, str) and c.value.startswith('Version ')]
        if len(vers) != 1 or not rx_version.match(vers[0]):
            R.fallo('%s: línea de versión D18 %r' % (en, vers))
        else:
            n['versiones'] += 1
        p = wb.properties
        esp = {'title': mapas.TITULOS[en], 'subject': A.SUBJECT, 'keywords': A.KEYWORDS, 'description': A.DESCRIPTION,
               'category': A.CATEGORY, 'creator': A.CREATOR}
        for k, v in esp.items():
            if getattr(p, k) != v:
                R.fallo('%s: docProps %s %r ≠ %r' % (en, k, getattr(p, k), v))
        if ins['B2'].value != mapas.titulo_interior(es):
            R.fallo('%s: Instructions!B2 %r ≠ título interior D5 %r' % (en, ins['B2'].value, mapas.titulo_interior(es)))
    # revisión final: rótulos con ajuste de texto y alto de fila suficiente (mapas.AJUSTE_TEXTO)
    for (corto, h_es, rg) in mapas.AJUSTE_TEXTO:
        ws = cargar(fichero_en(carpeta, corto))[mapas.HOJAS[h_es]]
        for coord in A.rango(rg):
            c = ws[coord]
            if c.value is None:
                continue
            n['ajuste_texto'] += 1
            alto = A.alto_necesario(ws, c.row)
            if not (c.alignment and c.alignment.wrap_text) or (alto > 15 and (ws.row_dimensions[c.row].height or 0) < alto):
                R.fallo('%s %s!%s: sin ajuste de texto o fila baja (alto %s < %s)' % (corto, ws.title, coord,
                                                                                   ws.row_dimensions[c.row].height, alto))
    T0 = censo['totales']
    n_pies = sum(len(hd['cabeceras_pies']) for i in censo['libros'].values() for hd in i['hojas_detalle'].values())
    if n['hojas'] != T0['hojas'] or n['pies'] != n_pies:
        R.fallo('hojas %d / pies %d ≠ censo %d / %d' % (n['hojas'], n['pies'], T0['hojas'], n_pies))
    if n['versiones'] != len(mapas.FICHEROS):
        R.fallo('versiones %d ≠ %d' % (n['versiones'], len(mapas.FICHEROS)))
    for k, v in sorted(n.items()):
        R.dato(k, v)


# ==========================================================================
# G4 — claves
# ==========================================================================
TABLAS_TIPOS = [   # (libro, hoja ES, rango de la tabla de tipos, hoja ES y celda de la DV que la alimenta)
    ('03', 'Ratio Coste Laboral', 'A18:A27', 'B3'),
    ('B02', 'Ratios por Tipo', 'A5:A14', None),
]


def gateG4(R, carpeta, datos):
    R.gate('G4 claves')
    n = Counter()
    listas_tipos = {}
    for es, en, corto, pes, pen in libros(carpeta):
        wb = cargar(pen)
        for ws in wb.worksheets:
            lugar = '%s:%s' % (en, ws.title)
            for dv in ws.data_validations.dataValidation:
                f1 = dv.formula1 or ''
                if dv.type == 'list' and f1.startswith('"'):
                    items = f1.strip('"').split(',')
                    n['dv_listas'] += 1
                    for x in items:
                        if mapas.necesita_clave(x) and x not in CLAVES_EN:
                            R.fallo('%s: DV %s ítem %r ∉ CLAVES' % (lugar, dv.sqref, x))
                    if any(x in {'Casual dining', 'Fine dining'} for x in items):
                        listas_tipos[corto] = set(items)
                    for coord in celdas(dv.sqref):
                        c = ws[coord]
                        if c.value is None or c.data_type == 'f':
                            continue
                        n['valores_con_dv_lista'] += 1
                        v = c.value
                        txt = ('%g' % v) if isinstance(v, (int, float)) and not isinstance(v, bool) else str(v)
                        if txt not in items:
                            R.fallo('%s!%s: %r fuera de su DV %s' % (lugar, coord, c.value, f1[:60]))
                elif dv.type in ('decimal', 'whole', 'time', 'date') and f1 and dv.formula2:
                    try:
                        lo, hi = float(f1), float(dv.formula2)
                    except ValueError:
                        continue
                    for coord in celdas(dv.sqref):
                        v = ws[coord].value
                        if v is None or ws[coord].data_type == 'f':
                            continue
                        n['valores_con_dv_numerica'] += 1
                        s = serial(v)
                        if s is None or not lo - 1e-9 <= s <= hi + 1e-9:
                            R.fallo('%s!%s: %r fuera de su DV %s-%s' % (lugar, coord, v, f1, dv.formula2))
    # tablas de tipos de negocio (claves de VLOOKUP) = la lista de su DV
    for corto, h_es, sq, _ in TABLAS_TIPOS:
        ws = cargar(fichero_en(carpeta, corto))[mapas.HOJAS[h_es]]
        tabla = {ws[c].value for c in A.rango(sq)}
        n['tipos_' + corto] = len(tabla)
        if tabla != listas_tipos.get(corto):
            R.fallo('%s %s!%s: tipos %r ≠ su DV %r' % (corto, ws.title, sq, sorted(map(str, tabla)),
                                                      sorted(listas_tipos.get(corto) or [])))
    for k, v in sorted(n.items()):
        R.dato(k, v)


# ==========================================================================
# G5 — CF
# ==========================================================================
def gateG5(R, carpeta, datos):
    R.gate('G5 CF')
    n = Counter()
    for es, en, corto, pes, pen in libros(carpeta):
        wes, wen = cargar(pes), cargar(pen)
        for h_es in wes.sheetnames:
            ws_es, ws_en = wes[h_es], wen[mapas.HOJAS[h_es]]
            tok_es, tok_en = {}, {}
            for ws, tok in ((ws_es, tok_es), (ws_en, tok_en)):
                for cf in ws.conditional_formatting:
                    for r in cf.rules:
                        s = tok.setdefault(str(cf.sqref), set())
                        for f in r.formula or []:
                            s.update(x for x in lits(f) if x)
                        if r.text:
                            s.add(r.text)
            for rg in sorted(set(tok_es) | set(tok_en)):
                e, m = tok_es.get(rg, set()), tok_en.get(rg, set())
                n['rangos'] += 1
                n['tokens'] += len(m)
                if len(e) != len(m) or len({x.lower() for x in m}) != len(m):
                    R.fallo('%s:%s CF %s: %d tokens ES → %d EN (colisión) %r' % (en, ws_en.title, rg, len(e), len(m), sorted(m)))
                fuera = sorted(x for x in m if mapas.necesita_clave(x) and x not in CLAVES_EN)
                if fuera:
                    R.fallo('%s:%s CF %s: tokens ∉ CLAVES %r' % (en, ws_en.title, rg, fuera))
                for a in m:
                    for b in m:
                        if a != b and a.lower() in b.lower():
                            esa = [k for k, v in CLAVES.items() if v == a]
                            esb = [k for k, v in CLAVES.items() if v == b]
                            if not any(x.lower() in y.lower() for x in esa for y in esb):
                                R.fallo('%s:%s CF %s: token %r contenido en %r (colisión SEARCH)' % (en, ws_en.title, rg, a, b))
    for e in mapas.cruzar_censo(os.path.join(AQUI, 'censo_es.json')):
        R.fallo('mapas.cruzar_censo: ' + e)
    for k, v in sorted(n.items()):
        R.dato(k, v)


# ==========================================================================
# G6 — cálculo
# ==========================================================================
def celdas_sin_v(path):
    fuera = set()
    with zipfile.ZipFile(path) as z:
        wbx = z.read('xl/workbook.xml').decode('utf-8')
        rels = z.read('xl/_rels/workbook.xml.rels').decode('utf-8')
        destino = {}
        for tag in re.findall(r'<Relationship\b[^>]*>', rels):
            i, t = re.search(r'\bId="([^"]+)"', tag), re.search(r'\bTarget="([^"]+)"', tag)
            if i and t:
                destino[i.group(1)] = t.group(1)
        for tag in re.findall(r'<sheet\b[^>]*>', wbx):
            nombre = html.unescape(re.search(r'\bname="([^"]*)"', tag).group(1))
            rid = re.search(r'\br:id="([^"]+)"', tag).group(1)
            t = destino[rid]
            parte = t.lstrip('/') if t.startswith('/') else 'xl/' + t
            xml = z.read(parte).decode('utf-8')
            for m in re.finditer(r'<c\b([^>]*?)(?:/>|>(.*?)</c>)', xml, re.S):
                dentro = m.group(2) or ''
                if '<f' in dentro and not re.search(r'<v>[^<]', dentro) and '<v></v>' not in dentro \
                        and '<v/>' not in dentro and '<v />' not in dentro:
                    fuera.add((nombre, re.search(r'\br="([A-Z]+\d+)"', m.group(1)).group(1)))
    return fuera


def num(v):
    return isinstance(v, (int, float)) and not isinstance(v, bool)


# Cifras que la SPEC cambia a propósito (D8, D20): (libro, hoja EN, celda) → valor EN esperado en la caché
ESPERADOS = OrderedDict([
    (('01', 'Shifts', 'E5'), 8), (('01', 'Shifts', 'E6'), 8), (('01', 'Shifts', 'E7'), 8), (('01', 'Shifts', 'E8'), 9),
    (('01', 'Shifts', 'E9'), 16), (('01', 'Shifts', 'E10'), 0), (('01', 'Shifts', 'E11'), 0), (('01', 'Shifts', 'E12'), 0),
    (('03', 'Staffing Forecast', 'B21'), 7), (('03', 'Staffing Forecast', 'B22'), 23356.67),
    (('B02', 'Staffing Calculator', 'B29'), 7), (('B02', 'Staffing Calculator', 'B38'), 23356.67),
    (('B02', 'Staffing Calculator', 'B39'), 280280.04), (('B02', 'Staffing Calculator', 'B40'), 72800),
    (('B02', 'Staffing Calculator', 'B41'), 0.3208), (('B02', 'Staffing Calculator', 'D41'), 0.33),
    (('B02', 'Staffing Calculator', 'B43'), CLAVES['🟢 EXCELENTE (por debajo de tu objetivo)']),
])
# D20: la frase de ejemplo de B02 Instructions!B16 lleva las cifras EN (y el veredicto)
D20_FRASE = ('B02', 'Instructions', 'B16', ('72,800', '23,356.67', '32.1', 'EXCELLENT'))
# D10: casos FLSA (SPEC §8 G6) en copias del 02 — (días, horas por día, F3) → horas extra totales
LUNES_PRUEBA = datetime.datetime(2027, 1, 4)      # semana laboral lunes 4-ene (B3 = 2, = rejilla del 01) → lunes a sábado
MIERCOLES_PRUEBA = datetime.datetime(2027, 1, 6)  # mié-dom (5 × 9 = 45 → 5 h) + lunes 11 en semana nueva (9 → 0) = 5 h;
#                                                   con la semana en domingo (B3 = 1) daría 0 h: prueba que B3 = 2 manda
CASOS_FLSA = [(6, 9, None, 14, LUNES_PRUEBA), (6, 9, 8, 14, LUNES_PRUEBA), (4, 10, None, 0, LUNES_PRUEBA),
              (4, 10, 8, 8, LUNES_PRUEBA), (6, 9, None, 5, MIERCOLES_PRUEBA)]


def caso_flsa(carpeta, dias, horas, f3, inicio=LUNES_PRUEBA):
    """Copia el 02, escribe un empleado con `dias` turnos de `horas` h (8:00 AM →), pone F3 y recalcula con
    inject_cache. Devuelve la suma de «Time Log»!H y la lista por fila."""
    tmp = tempfile.mkdtemp(prefix='staff-flsa-')
    try:
        src = fichero_en(carpeta, '02')
        p = os.path.join(tmp, os.path.basename(src))
        shutil.copy2(src, p)
        wb = openpyxl.load_workbook(p)
        ws = wb[mapas.HOJAS['Registro Horas']]
        ws['F3'].value = f3
        for i in range(dias):
            r = 5 + i
            ws['A%d' % r].value = 'FLSA test'
            ws['B%d' % r].value = inicio + datetime.timedelta(days=i)
            ws['C%d' % r].value = datetime.time(8, 0)
            ws['D%d' % r].value = datetime.time(8 + horas, 0)
            ws['E%d' % r].value = 0
        wb.save(p)
        r = subprocess.run([sys.executable, '-W', 'ignore', A.INJECT, p], stdout=subprocess.PIPE,
                           stderr=subprocess.PIPE, universal_newlines=True)
        if r.returncode or 'fallos_pycel=0' not in r.stdout:
            return None, 'inject_cache: %s %s' % (r.stdout.strip()[-120:], r.stderr[-200:])
        wv = openpyxl.load_workbook(p, data_only=True)[mapas.HOJAS['Registro Horas']]
        hs = [wv['H%d' % (5 + i)].value for i in range(dias)]
        fs = [wv['F%d' % (5 + i)].value for i in range(dias)]
        if any(not num(x) for x in hs + fs):
            return None, 'caché vacía o no numérica: F %r · H %r' % (fs, hs)
        return round(sum(hs), 2), 'F %r · H %r' % (fs, hs)
    finally:
        shutil.rmtree(tmp)


def gateG6(R, carpeta, datos, flsa=True):
    R.gate('G6 cálculo')
    n = Counter()
    cambian = {(c, h, coord) for (c, h, coord) in ESPERADOS}
    for es, en, corto, pes, pen in libros(carpeta):
        wf, wv = cargar(pen), cargar(pen, data_only=True)
        wes_v = cargar(pes, data_only=True)
        wes_f = cargar(pes)
        sin_v = celdas_sin_v(pen)
        for h_es in wes_f.sheetnames:
            h_en = mapas.HOJAS[h_es]
            ws, ws_v, we_v = wf[h_en], wv[h_en], wes_v[h_es]
            for c in ws._cells.values():
                if c.data_type != 'f':
                    continue
                n['formulas'] += 1
                if (ws.title, c.coordinate) in sin_v:
                    R.fallo('%s:%s!%s sin caché (<f> sin <v>)' % (en, ws.title, c.coordinate))
                    continue
                v = ws_v[c.coordinate].value
                ve = we_v[c.coordinate].value
                if isinstance(v, str) and v in ERRORES:
                    R.fallo('%s:%s!%s = %s' % (en, ws.title, c.coordinate, v))
                    continue
                n['con_valor' if v is not None else 'vacias_legitimas'] += 1
                if (corto, ws.title, c.coordinate) in cambian:
                    continue
                if (ve is None) != (v is None):
                    R.fallo('%s:%s!%s caché %r (ES %r)' % (en, ws.title, c.coordinate, v, ve))
                elif isinstance(ve, str) and ve in CLAVES:
                    n['veredictos'] += 1
                    if v != CLAVES[ve]:
                        R.fallo('%s:%s!%s veredicto %r ≠ CLAVES[%r] = %r' % (en, ws.title, c.coordinate, v, ve, CLAVES[ve]))
                elif isinstance(v, str):
                    n['textos_calculados'] += 1
                    for motivo, ctx in restos_espanol(v):
                        R.fallo('%s:%s!%s texto calculado con %s · …%s…' % (en, ws.title, c.coordinate, motivo, ctx))
                    if '€' in v:
                        R.fallo('%s:%s!%s texto calculado con €: %r' % (en, ws.title, c.coordinate, v))
                elif num(ve):
                    n['numeros'] += 1
                    if not (num(v) and abs(v - ve) < TOL):
                        R.fallo('%s:%s!%s número %r ≠ ES %r' % (en, ws.title, c.coordinate, v, ve))
                elif isinstance(ve, (datetime.date, datetime.datetime)):
                    n['fechas'] += 1
                    if not isinstance(v, (datetime.date, datetime.datetime)):
                        R.fallo('%s:%s!%s fecha %r (ES %r)' % (en, ws.title, c.coordinate, v, ve))
    # cifras que cambian por la SPEC
    for (corto, hoja, coord), e in ESPERADOS.items():
        v = cargar(fichero_en(carpeta, corto), True)[hoja][coord].value
        n['esperados'] += 1
        ok = (num(v) and num(e) and abs(v - e) < (0.00051 if abs(e) < 1 else 0.006)) or v == e
        if not ok:
            R.fallo('%s %s!%s = %r ≠ SPEC %r' % (corto, hoja, coord, v, e))
    corto, hoja, coord, agujas = D20_FRASE
    t = cargar(fichero_en(carpeta, corto))[hoja][coord].value or ''
    faltan = [a for a in agujas if a not in t]
    if faltan:
        R.fallo('D20 %s %s!%s no lleva %r: %r' % (corto, hoja, coord, faltan, t[:140]))
    # 05: la semana 1 empieza en domingo y las 52 siguientes van de 7 en 7
    w5 = cargar(fichero_en(carpeta, '05'), True)[mapas.HOJAS['Calendario Anual']]
    b5, c5 = w5['B5'].value, w5['C5'].value
    if not isinstance(b5, datetime.datetime) or b5.weekday() != 6 or not isinstance(c5, datetime.datetime) \
            or (c5 - b5).days != 7:
        R.fallo('05 Annual Calendar B5/C5 = %r / %r (semana 1 en domingo, +7)' % (b5, c5))
    # D10: casos FLSA
    if flsa:
        for dias, horas, f3, esperado, inicio in CASOS_FLSA:
            tot, det = caso_flsa(carpeta, dias, horas, f3, inicio)
            n['casos_flsa'] += 1
            R.dato('flsa_%dx%d_F3=%s_%s' % (dias, horas, f3, inicio.strftime('%a')), tot)
            if tot is None or abs(tot - esperado) > 0.001:
                R.fallo('FLSA %d × %d h con F3=%s → %r ≠ %s h (%s)' % (dias, horas, f3, tot, esperado, det))
    for k, v in sorted(n.items()):
        R.dato(k, v)


# ==========================================================================
# G7 — reglas (SPEC §3 = mapas.REGLAS_US)
# ==========================================================================
def gateG7(R, carpeta, datos):
    R.gate('G7 reglas')
    n = Counter()
    RU = {k: v[0] for k, v in mapas.REGLAS_US.items()}
    # 1. VALORES_EN en los xlsx
    for (corto, h_es, ref), v in mapas.VALORES_EN.items():
        ws = cargar(fichero_en(carpeta, corto))[mapas.HOJAS[h_es]]
        coords = list(A.rango(ref))
        wes = cargar(os.path.join(ORIGEN, [e for e in mapas.LIBROS_ES if A.extraer_textos.corto(e) == corto][0]))[h_es]
        for coord in coords:
            if len(coords) > 1 and wes[coord].value is None:
                continue
            got = ws[coord].value
            ok = got == v or (isinstance(v, datetime.date) and isinstance(got, datetime.datetime) and got.date() == v) \
                or (num(got) and num(v) and abs(got - v) < 1e-9)
            n['valores'] += 1
            if not ok:
                R.fallo('VALORES_EN %s %s!%s = %r ≠ %r' % (corto, ws.title, coord, got, v))
    # 2. las cifras de VALORES_EN / PARAMETROS que vienen de la norma coinciden con SPEC §3
    V = mapas.VALORES_EN
    P = {(c, h, k): val for (c, h, k), (et, val) in mapas.PARAMETROS.items()}
    comprobaciones = [
        ('ot_semanal 40 h', P[('02', 'Registro Horas', 'D3')] == 40 and '40 h' in RU['ot_semanal']),
        ('ot_semanal 1.5×', V[('02', 'Resumen Mensual', 'D3')] == 1.5 and '1.5' in RU['ot_semanal']),
        ('F3 vacío = federal', P[('02', 'Registro Horas', 'F3')] is None),
        ('semana laboral 2 = lunes (= rejilla lunes-domingo del 01)', P[('02', 'Registro Horas', 'B3')] == 2),
        ('turno_largo 10 h', V[('01', 'Turnos', 'B2')] == 10 and RU['turno_largo'].startswith('10 h')),
        ('descanso_turnos 10 h', V[('01', 'Turnos', 'B3')] == 10 and RU['descanso_turnos'].startswith('10 h')),
        ('coste_empresa 10 %', V[('03', 'Nóminas', 'C2')] == 0.10 == V[('03', 'Previsión por Servicio', 'B16')]
         == V[('B02', 'Calculadora', 'B11')] and RU['coste_empresa'] == '10 %'),
        ('periodos 26', V[('03', 'Nóminas', 'D5:D34')] == 26 == V[('03', 'Previsión por Servicio', 'B15')]
         == V[('B02', 'Calculadora', 'B10')] and RU['periodos_pago'].startswith('52/26/24/12, 26')),
        ('semanas 50', V[('03', 'Nóminas', 'F2')] == 50 and RU['semanas'] == '50'),
        ('pto 14', V[('05', 'Saldo Vacaciones', 'B2')] == 14 and RU['pto'].startswith('14')),
        ('temp 41 °F', V[('B01', 'Briefing', 'D56')] == 41 == V[('B01', 'Briefing', 'D58')] and '41 °F' in RU['temp']),
        ('temp 135 °F', V[('B01', 'Briefing', 'C59')] == 135 == V[('B01', 'Briefing', 'C60')] and '135 °F' in RU['temp']),
        ('congelador ≤ 0 °F', V[('B01', 'Briefing', 'D57')] == 0),
        ('semana 1 en domingo', V[('05', 'Calendario Anual', 'B5')].weekday() == 6),
    ]
    for nombre, ok in comprobaciones:
        n['reglas'] += 1
        if not ok:
            R.fallo('regla %s: mapas ≠ SPEC §3' % nombre)
    # 3. listas especiales D11 = periodos de pago de SPEC §3
    for k, (es_l, en_l) in mapas.DV_LISTAS_ESPECIALES.items():
        if en_l.strip('"').replace(',', '/') != RU['periodos_pago'].split(',')[0]:
            R.fallo('DV_LISTAS_ESPECIALES %r = %s ≠ SPEC §3 %s' % (k, en_l, RU['periodos_pago']))
    # 4. PARAMETROS escritos en «Time Log» (fila 3) y verdes/desbloqueados
    wl = cargar(fichero_en(carpeta, '02'))[mapas.HOJAS['Registro Horas']]
    for (c, h, coord), (et, val) in mapas.PARAMETROS.items():
        got = wl[coord].value
        n['parametros'] += 1
        if got != (et if et is not None else val):
            R.fallo('PARAMETROS 02 Time Log!%s = %r ≠ %r' % (coord, got, et if et is not None else val))
    # 5. DV_NUEVAS: los mensajes citan las cifras de la regla
    dvn = mapas.DV_NUEVAS
    if '40' not in dvn[('02', 'Registro Horas', 'D3')][4] or 'California: 8' not in dvn[('02', 'Registro Horas', 'F3')][4]:
        R.fallo('DV_NUEVAS: los mensajes de D3/F3 no citan 40 h / California 8 h')
    # 6. FORMULAS_EN: la H de «Time Log» usa los tres parámetros (semana laboral, umbral semanal, umbral diario)
    for c in wl._cells.values():
        if c.data_type == 'f' and c.column_letter == 'H' and c.row >= 5:
            n['h_flsa'] += 1
            for frag in ('$B$3', '$D$3', '$F$3', 'WEEKDAY'):
                if frag not in c.value:
                    R.fallo('02 Time Log!%s sin %s' % (c.coordinate, frag))
                    break
    if n['h_flsa'] != 300:
        R.fallo('02 Time Log!H: %d fórmulas FLSA ≠ 300' % n['h_flsa'])
    # 7. Shifts: TURNOS_EN (mismos horarios que el ES) y los códigos = DV de la rejilla
    ws = cargar(fichero_en(carpeta, '01'))['Shifts']
    for i, (cod, desc, ini, fin, pausa) in enumerate(mapas.TURNOS_EN):
        r = A.TURNOS_FILA0 + i
        got = (ws['A%d' % r].value, ws['B%d' % r].value, serial(ws['C%d' % r].value), serial(ws['D%d' % r].value),
               ws['F%d' % r].value)
        esp = (cod, desc, serial(A.hora(ini)), serial(A.hora(fin)), pausa)
        n['turnos'] += 1
        if got[:2] != esp[:2] or got[4] != esp[4] or abs((got[2] or 0) - esp[2]) > 1e-9 or abs((got[3] or 0) - esp[3]) > 1e-9:
            R.fallo('Shifts fila %d = %r ≠ TURNOS_EN %r' % (r, got, esp))
    for k, v in sorted(n.items()):
        R.dato(k, v)


# ==========================================================================
# G8 — censo-entregables y gate-no-latinos
# ==========================================================================
def gateG8(R, carpeta, datos):
    R.gate('G8 censo + no-latinos')
    real = os.path.abspath(carpeta) == os.path.abspath(A.DESTINO_REAL)
    cmds = [[sys.executable, CENSO_SCRIPT, '--only', mapas.SLUG if real else carpeta, '--fail'] + ([] if real else ['--letter']),
            [sys.executable, NO_LATINOS, '--only', mapas.SLUG if real else carpeta]]
    for cmd in cmds:
        r = subprocess.run(cmd, stdout=subprocess.PIPE, stderr=subprocess.STDOUT, universal_newlines=True, cwd=REPO)
        lineas = [x for x in r.stdout.strip().splitlines() if x.strip()]
        nombre = os.path.basename(cmd[1])
        R.dato(nombre, lineas[-1] if lineas else '')
        if r.returncode:
            R.fallo('%s exit %d: %s' % (nombre, r.returncode, ' | '.join(lineas[-6:])))


# ==========================================================================
# Autotest: cada gate tiene que cazar un defecto inyectado en una COPIA
# ==========================================================================
def _mutar(carpeta, fname, fn):
    p = os.path.join(carpeta, fname)
    wb = openpyxl.load_workbook(p)
    fn(wb)
    wb.save(p)


def _set(ws, coord, v):
    ws[coord].value = v


def _cf_colision(wb):
    ws = wb['Monthly OT Summary']
    for cf in ws.conditional_formatting:
        for r in cf.rules:
            r.formula = [f.replace('"near budget"', '"OVER BUDGET"') for f in r.formula]
            if r.text == 'near budget':
                r.text = 'OVER BUDGET'


def _lock(wb):
    from openpyxl.styles import Protection
    wb['Time Log']['B3'].protection = Protection(locked=True)


def _dv_titulo_id(wb):
    for dv in wb['Monthly OT Summary'].data_validations.dataValidation:
        if str(dv.sqref) == 'B3':
            dv.promptTitle = 'kitgp-v2 · hourly_rate'


def _sin_ajuste(wb):
    from openpyxl.styles import Alignment
    wb['Onboarding Checklist']['B14'].alignment = Alignment(wrap_text=False)


def _formula_shifts(wb):
    ws = wb['Shifts']
    ws['E5'].value = ws['E5'].value.replace('*24', '')


MUTACIONES = [
    ('G1', '01-restaurant-schedule-template.xlsx', _formula_shifts, 'fórmula EN'),
    ('G1', '02-overtime-tracker.xlsx', _lock, 'desbloqueadas'),
    ('G2', '04-new-hire-onboarding-checklist.xlsx', lambda wb: _set(wb['Instructions'], 'B3', 'Revisa cada tarea con el empleado'), 'ES'),
    ('G2', '03-labor-cost-calculator.xlsx', lambda wb: _set(wb['Instructions'], 'B3', 'Average pay: 1,400 € per period'), 'moneda'),
    ('G3', '06-employee-performance-review.xlsx', lambda wb: setattr(wb['Review History'].page_setup, 'paperSize', 9), 'paperSize'),
    ('G3', '05-pto-vacation-planner.xlsx', lambda wb: setattr(wb['Annual Calendar']['B5'], 'number_format', 'dd/mm'), 'formato'),
    ('G2', '02-overtime-tracker.xlsx', _dv_titulo_id, 'ID interno'),
    ('G2', '03-labor-cost-calculator.xlsx', lambda wb: _set(wb['Instructions'], 'B9', 'See file 02 (Overtime Tracker '
                                                         '& Time Log (FLSA 40-Hour Week)).'), 'cita anidada'),
    ('G3', '04-new-hire-onboarding-checklist.xlsx', _sin_ajuste, 'ajuste de texto'),
    ('G4', '01-restaurant-schedule-template.xlsx', lambda wb: _set(wb['Weekly Schedule'], 'B6', 'M'), 'fuera de su DV'),
    ('G5', '02-overtime-tracker.xlsx', _cf_colision, 'colisión'),
    ('G6', 'BONUS-02-restaurant-staffing-calculator.xlsx', lambda wb: _set(wb['Staffing Calculator'], 'B15', 25), 'SPEC'),
    ('G7', '01-restaurant-schedule-template.xlsx', lambda wb: _set(wb['Shifts'], 'B3', 12), 'VALORES_EN'),
]


def autotest(carpeta, mes, datos):
    ok = True
    base = ejecutar(carpeta, mes, ['G1', 'G2', 'G3', 'G4', 'G5', 'G6', 'G7'], datos, silencioso=True, flsa=False)
    sucios = [g for g, v in base.gates.items() if v['fallos']]
    if sucios:
        print('AUTOTEST: la carpeta base no está limpia (%s): arregla eso primero' % sucios)
        return False
    for gate, fname, fn, aguja in MUTACIONES:
        tmp = tempfile.mkdtemp(prefix='staff-autotest-')
        try:
            for f in mapas.FICHEROS.values():
                shutil.copy2(os.path.join(carpeta, f), tmp)
            _mutar(tmp, fname, fn)
            if gate in ('G6',):
                subprocess.run([sys.executable, '-W', 'ignore', A.INJECT, os.path.join(tmp, fname)],
                               stdout=subprocess.PIPE, stderr=subprocess.PIPE)
            R = ejecutar(tmp, mes, [gate], datos, silencioso=True, flsa=False)
            fallos = [f for g in R.gates.values() for f in g['fallos']]
            cazado = any(aguja.lower() in f.lower() for f in fallos)
            ok &= cazado
            print('  autotest %s %-46s %s' % (gate, fname, 'CAZADO' if cazado else 'NO CAZADO %r' % fallos[:2]))
        finally:
            shutil.rmtree(tmp)
    return ok


# ==========================================================================
# main
# ==========================================================================
GATES = OrderedDict([('G1', gateG1), ('G2', gateG2), ('G3', gateG3), ('G4', gateG4), ('G5', gateG5),
                     ('G6', gateG6), ('G7', gateG7), ('G8', gateG8)])


def cargar_datos():
    return {'censo': json.load(open(os.path.join(AQUI, 'censo_es.json'), encoding='utf-8'))}


def ejecutar(carpeta, mes, solo, datos, silencioso=False, flsa=True):
    R = Resultado()
    for nombre, fn in GATES.items():
        if solo and nombre not in solo:
            continue
        A.termica()
        if nombre == 'G3':
            fn(R, carpeta, datos, mes)
        elif nombre == 'G6':
            fn(R, carpeta, datos, flsa=flsa)
        else:
            fn(R, carpeta, datos)
        g = R.gates[R.actual]
        if not silencioso:
            estado = 'OK' if not g['fallos'] else 'FALLA (%d)' % len(g['fallos'])
            resumen = ' · '.join('%s=%s' % (k, v) for k, v in list(g['datos'].items())[:14] if not isinstance(v, list))
            print('%-24s %s  %s' % (R.actual, estado, resumen[:300]))
            for f in g['fallos'][:int(os.environ.get('STAFF_MAX_FALLOS', '15'))]:
                print('    - ' + f[:260])
    return R


def main():
    ap = argparse.ArgumentParser(description='Gates F2 del Restaurant Staff Scheduling Kit Pro (EN)')
    ap.add_argument('--dir', default=A.DESTINO_REAL, help='carpeta con los 9 xlsx EN (por defecto dl/<slug>)')
    ap.add_argument('--mes', default=None, help='mes de la línea de versión (por defecto, el actual)')
    ap.add_argument('--solo', default=None, help='G1,G2,… (por defecto, todos)')
    ap.add_argument('--json', default=None)
    ap.add_argument('--autotest', action='store_true')
    args = ap.parse_args()
    mes = args.mes or A.mes_por_defecto()
    datos = cargar_datos()
    carpeta = os.path.abspath(args.dir)
    if args.autotest:
        ok = autotest(carpeta, mes, datos)
        print('AUTOTEST: %s' % ('OK' if ok else 'FALLA'))
        sys.exit(0 if ok else 1)
    solo = [s.strip().upper() for s in args.solo.split(',')] if args.solo else None
    R = ejecutar(carpeta, mes, solo, datos)
    fallos = sum(len(g['fallos']) for g in R.gates.values())
    if args.json:
        with open(args.json, 'w', encoding='utf-8') as fh:
            json.dump(R.gates, fh, ensure_ascii=False, indent=1, default=str)
    print('\n%s: %d gates, %d fallos (%s)' % ('VERDE' if not fallos else 'ROJO', len(R.gates), fallos, carpeta))
    sys.exit(1 if fallos else 0)


if __name__ == '__main__':
    main()
