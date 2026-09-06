#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
gen_auditoria-interna-cocina.py — libro 7 de «Manual del Chef Ejecutivo»
(SPEC §2.2 fila 7).

Hojas: Instrucciones · Auditoría · Resumen por Área · Histórico · Histórico por
Unidad · Estado Normativo.

QUÉ DECIDE ESTE LIBRO
---------------------
Puntuar la disciplina de cocina de forma repetible, ver si mejora visita a
visita y comparar unas cocinas con otras. EXCLUYE a propósito limpieza, plagas y
temperaturas de cámara: eso es el Pack APPCC, y mezclarlo produce un checklist
que no sirve ni para una cosa ni para la otra.

DECISIONES TÉCNICAS
-------------------
* La puntuación es una MEDIA PONDERADA, `SUMPRODUCT(peso;puntuación)/SUM(peso)`,
  tal cual: no es la media aritmética de los 50 puntos. Un punto de peso 3
  («los gramajes servidos coinciden con la ficha») no vale lo mismo que uno de
  peso 1 («las estanterías de seco están rotuladas»).
* Los 50 puntos van AGRUPADOS por área en bloques contiguos de diez, y el
  resumen por área usa el rango exacto de cada bloque. Así no hace falta ni un
  `SUMIF` con criterio de texto ni un `SUMPRODUCT` con comparaciones: menos
  cosas que se rompan cuando alguien reordene una fila.
* El área de prevención de riesgos va SEMBRADA con base legal verificada
  (CE-21 a CE-27), cada punto con su norma, su enlace y la nota «Verificado el
  06-09-2026 · norma · URL» leída del JSON de research. Es lo que le da
  herramienta al bloque de PRL sin añadir un octavo libro.
* El semáforo de la auditoría NO usa ningún benchmark del sector, porque no
  existe ninguno publicado en español: son dos celdas verdes (aviso ámbar y
  crítico) sobre la escala del libro, y la comparación que de verdad manda es
  contra TU visita anterior. En el ejemplo la media global sube y sin embargo un
  área EMPEORA: eso es justo lo que una media global esconde.
* `Histórico por Unidad` se sirve con la unidad de referencia calculada y las
  demás en blanco: `datos_ejemplo.py` no trae auditorías de las otras cocinas
  del grupo y aquí no se inventa ninguna puntuación.
* Cero constantes dentro de una fórmula (salvo 0, 1 y los índices de
  INDEX/MATCH), `IFERROR(...,"")` en todo cociente, «sin dato» = `""` nunca 0,
  semáforos con `ISNUMBER`.
* Prohibidas `INDIRECT`, `COUNTA`, `PMT`, `OFFSET`, `XLOOKUP`, `LET`, `LAMBDA`,
  `RANK`, `NETWORKDAYS` y las referencias a otros ficheros: cero usos.
* Textos 100 % WinAnsi (cp1252): ni un carácter fuera.

Salida fija: build/auditoria-interna-cocina.xlsx + su mapa de celdas.
"""
import json
import os
import sys

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.page import PageMargins
from openpyxl.worksheet.properties import PageSetupProperties

AQUI = os.path.dirname(os.path.abspath(__file__))
RAIZ = os.path.normpath(os.path.join(AQUI, '..'))
sys.path.insert(0, os.path.join(RAIZ, 'guias-v2_0'))
sys.path.insert(0, AQUI)

import motor                                                   # noqa: E402
import datos_ejemplo as D                                      # noqa: E402

motor.CTX['producto'] = 'manual-chef-ejecutivo'

PRODUCTO = 'Manual del Chef Ejecutivo'
SUBTITULO = 'AI Chef Pro · aichef.pro — ' + PRODUCTO
TITULO = 'Auditoría interna de cocina'
NOMBRE = 'auditoria-interna-cocina'

FMT_EUR = motor.FMT_EUR
FMT_PCT = motor.FMT_PCT
FMT_ENT = motor.FMT_ENT
FMT_FECHA = motor.FMT_FECHA
FMT_DEC = '#,##0.00'

GRIS = 'F2F2F2'
CREMA = 'FFF6DC'
ORO = 'FFD700'
CABECERA = '2D2D2D'

FECHA_VERIF = '06-09-2026'
JSON_SECTOR = os.path.join(RAIZ, 'auditorias', 'guias-v2-research-sector.json')

with open(JSON_SECTOR, encoding='utf-8') as _fh:
    SECTOR = dict((it['id'], it) for it in json.load(_fh)['datos']
                  if it.get('id'))


def verificado(idd):
    """Nota literal de la familia: «Verificado el <fecha> · norma · URL»."""
    if idd not in SECTOR:
        raise KeyError('id de research inexistente: ' + idd)
    it = SECTOR[idd]
    return ('Verificado el ' + FECHA_VERIF + ' · ' + it['fuente_titulo']
            + ' · ' + it['url'])


AREAS = list(D.AREAS_AUDITORIA_COCINA)
#: bloques CONTIGUOS de diez puntos por área, en el orden de la lista
BLOQUES = []
for _a in AREAS:
    _idx = [i for i, f in enumerate(D.AUDITORIA_COCINA) if f[1] == _a]
    BLOQUES.append((_idx[0], _idx[-1]))


def _nota_ponderada(puntuaciones):
    """SUMPRODUCT(peso;puntuación)/SUM(peso), igual que la fórmula de la
    hoja «Histórico». No es una media aritmética."""
    pesos = [f[3] for f in D.AUDITORIA_COCINA]
    return sum(p * s for p, s in zip(pesos, puntuaciones)) / sum(pesos)


#: A4 de la refutación xlsx (2026-09-06): «Tu estándar de grupo» venía en
#: blanco, así que «Histórico por Unidad» nunca demostraba las columnas
#: «Desviación contra tu estándar» y «Lectura». Se siembra con la nota
#: ponderada de la última visita propia (la del ejemplo), no un número suelto.
ESTANDAR_GRUPO_EJEMPLO = round(
    _nota_ponderada(D.AUDITORIAS_COCINA_HECHAS[-1][3]), 2)

# --- hoja «Auditoría» ----------------------------------------------------
A_FECHA, A_AUDITOR, A_UNIDAD = 7, 8, 9
A_ESCALA, A_AMBAR, A_ROJO = 10, 11, 12
A_MEDIA, A_CUMP, A_PUNT = 13, 14, 15
A_CAB = 18
A_INI = 19
A_FIN = A_INI + len(D.AUDITORIA_COCINA) - 1                      # 68

# --- hoja «Resumen por Área» ---------------------------------------------
S_COMP = 4                                   # visita del histórico a comparar
S_CAB = 5
S_INI = 6
S_FIN = S_INI + len(AREAS) - 1                                   # 10
S_TOT = S_FIN + 2                                                # 12

# --- hoja «Histórico» ----------------------------------------------------
H_PESOS = 4
H_CAB = 5
H_INI = 6
H_FIN = H_INI + 8 - 1                                            # 13

# --- hoja «Histórico por Unidad» -----------------------------------------
U_PESOS = 4
U_CAB = 5
U_INI = 6
U_FIN = U_INI + 6 - 1                                            # 11
U_MEDIA = U_FIN + 2                                              # 13
U_ESTANDAR = U_MEDIA + 1                                         # 14

# --- hoja «Estado Normativo» ---------------------------------------------
N_CORTE, N_SIM, N_HOY, N_DIAS, N_UMBRAL, N_AVISO = 5, 6, 7, 8, 9, 10
N_CAB = 12
N_INI = 13
N_FIN = N_INI + len(D.ESTADO_NORMATIVO) - 1                      # 27

AUD = "'Auditoría'"
RES = "'Resumen por Área'"
HIS = "'Histórico'"

LECT_MAL = 'EMPEORA respecto a la visita comparada'
LECT_BIEN = 'Mejora o se mantiene'
LECT_OBJ_MAL = 'Por debajo de tu objetivo'
LECT_OBJ_BIEN = 'En tu objetivo'
AVISO_VIEJO = 'REVISAR: ha pasado mucho desde la fecha de corte'
AVISO_NUEVO = 'Revisión reciente'
ESTADOS_REVISION = ['Sin comprobar', 'Sigue igual', 'Ha cambiado']

#: días tras los que el libro sugiere volver a comprobar las normas. Criterio
#: de la casa en celda verde, no una obligación de ninguna norma.
DIAS_AVISO_REVISION = 180


# --------------------------------------------------------------------------
def cabecera(ws, fila, columnas, altura=44):
    for letra, texto in columnas:
        c = ws[letra + str(fila)]
        c.value = texto
        c.fill = PatternFill('solid', fgColor=CABECERA)
        c.font = Font(bold=True, color='FFFFFF')
        c.alignment = Alignment(horizontal='center', vertical='center',
                                wrap_text=True)
    ws.row_dimensions[fila].height = altura


def pintar_cabecera(ws, fila, letras):
    for letra in letras:
        ws[letra + str(fila)].fill = PatternFill('solid', fgColor=CABECERA)


def seccion(ws, coord, texto):
    motor.val(ws, coord, texto, bold=True)
    ws[coord].font = Font(bold=True, size=12)


def fila_total(ws, fila, primera, ultima):
    for col in range(motor.column_index_from_string(primera),
                     motor.column_index_from_string(ultima) + 1):
        c = ws.cell(row=fila, column=col)
        c.fill = PatternFill('solid', fgColor=CREMA)
        c.font = Font(bold=True)


def cf_expresion(ws, rango, formula, bg, fg):
    """Regla de formato condicional que SE APILA sobre las anteriores.

    `motor.regla_expresion` purga las reglas del mismo rango antes de escribir
    (es lo correcto cuando hay una sola), así que para los tres tramos de un
    semáforo hace falta esta variante, que es la que usa la familia.
    """
    ws.conditional_formatting.add(
        rango, FormulaRule(formula=[formula], stopIfTrue=True,
                           font=Font(color=fg, bold=True),
                           fill=PatternFill(start_color=bg, end_color=bg,
                                            fill_type='solid')))


def pagina(ws, apaisado=True, titulos=None):
    ws.page_setup.paperSize = 9
    ws.page_setup.orientation = 'landscape' if apaisado else 'portrait'
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    if ws.sheet_properties.pageSetUpPr is None:
        ws.sheet_properties.pageSetUpPr = PageSetupProperties()
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_margins = PageMargins(left=0.59, right=0.59, top=0.59,
                                  bottom=0.59, header=0.3, footer=0.3)
    ws.oddFooter.center.text = 'AI Chef Pro · aichef.pro · Página &P de &N'
    ws.oddFooter.center.size = 8
    if titulos:
        ws.print_title_rows = titulos


def encabezar(ws, titulo, nota=None):
    motor.val(ws, 'A1', titulo)
    ws['A1'].font = Font(bold=True, size=16, color=ORO)
    ws.row_dimensions[1].height = 30
    motor.val(ws, 'A2', SUBTITULO)
    if nota:
        motor.val(ws, 'A3', nota)
        ws['A3'].font = Font(italic=True, size=9)


def anchos(ws, mapa):
    for letra, ancho in mapa.items():
        ws.column_dimensions[letra].width = ancho


def nota(ws, fila, texto, alto=None, col='A', wrap=False):
    motor.val(ws, col + str(fila), texto, wrap=wrap)
    if alto:
        ws.row_dimensions[fila].height = alto


def gris(ws, coord):
    ws[coord].fill = PatternFill('solid', fgColor=GRIS)


def dv_rango(ws, coords, ref, titulo, mensaje):
    """Desplegable contra un RANGO (convención de familia: nunca lista con
    comas). A1 de la refutación xlsx del 2026-09-06."""
    dv = DataValidation(type='list', formula1=ref, allow_blank=True,
                        showErrorMessage=True, errorTitle=titulo,
                        error=mensaje)
    ws.add_data_validation(dv)
    for c in coords:
        dv.add(c)
    return dv


#: Bloque de la lista del desplegable de «Estado Normativo» (A1 de la
#: refutación xlsx): al final de la hoja, con margen tras las tres notas.
N_LSEC = N_FIN + 6
N_REV_CAB = N_LSEC + 1
N_REV_INI = N_REV_CAB + 1
N_REV_FIN = N_REV_INI + len(ESTADOS_REVISION) - 1
REF_ESTADOS_REVISION = ("'Estado Normativo'!$A$%d:$A$%d"
                        % (N_REV_INI, N_REV_FIN))


# --------------------------------------------------------------------------
PASOS = [
    '1. Hoja «Auditoría»: escribe la fecha, quién audita y qué cocina se '
    'audita. La escala viene a 5 y los dos umbrales del semáforo (ámbar y '
    'crítico) son celdas verdes: son tuyos, no de ningún estándar.',
    '2. Recorre los 50 puntos y puntúa cada uno de 0 a la escala. El peso ya '
    'viene puesto de 1 a 3 y también es editable: si en tu casa la trazabilidad '
    'de las fichas pesa más que el orden del office, súbelo.',
    '3. La nota de la visita es una MEDIA PONDERADA, no la media de los 50 '
    'números: cada punto pesa lo que dice su columna. Está arriba, junto al '
    'cumplimiento sobre la escala y los puntos que llevas puntuados.',
    '4. Hoja «Resumen por Área»: los mismos 50 puntos repartidos en cinco '
    'áreas, cada una con su nota, su peso y su comparación contra la visita '
    'que elijas en la celda verde de arriba. La columna «Lectura» dice con '
    'palabras si el área empeora.',
    '5. Esa comparación por área es LO IMPORTANTE de este libro. En el ejemplo '
    'la nota global sube de 3,09 a 3,54 y aun así el área de mermas cae de '
    '3,67 a 2,58: una media global se traga eso y una comparación por área no.',
    '6. Hoja «Histórico»: una fila por visita. La primera es la visita ya '
    'cerrada, con sus notas por área escritas a mano; la segunda es la VISITA '
    'EN CURSO y se lee sola de las otras dos hojas. Cuando la cierres, copia '
    'esa fila y pégala como valores en la primera fila libre de abajo.',
    '7. Hoja «Histórico por Unidad»: la misma auditoría puntuada en varias '
    'cocinas. La unidad de referencia se calcula sola; las demás las rellenas '
    'tú cuando las audites. Debajo tienes la media del grupo y tu estándar, '
    'para ver quién se aparta y por dónde.',
    '8. Hoja «Estado Normativo»: qué dice hoy cada norma que toca esta '
    'auditoría, qué tiene que hacer el chef, la fecha en que se comprobó y el '
    'enlace para comprobarlo tú. La fecha de corte es editable y el libro te '
    'avisa cuando ha pasado demasiado tiempo desde ella.',
    '9. Lo que esta auditoría NO mira: limpieza, control de plagas y '
    'temperaturas de cámara. Eso son registros de APPCC y viven en el «Pack '
    'APPCC». Mezclarlos aquí daría un checklist que no sirve ni para una cosa '
    'ni para la otra.',
    '10. Cadencia: una visita AL CIERRE DE TRIMESTRE o de temporada. No es una '
    'herramienta de uso diario ni semanal: es la fotografía que compara una '
    'temporada con la anterior.',
]

NOTAS_LIBRO = [
    'ESTA AUDITORÍA MIDE DISCIPLINA DE COCINA, no higiene documental. Son '
    'cincuenta cosas que se ven en un servicio: si la mise en place está '
    'montada a la hora, si los gramajes coinciden con la ficha, si la merma se '
    'pesa o se estima, si el plato con alergia sale identificado.',
    'La nota ponderada es la única honesta. Con media aritmética, rotular las '
    'estanterías de seco (peso 1) sube lo mismo que pesar tres raciones contra '
    'la ficha (peso 3), y entonces la auditoría se puede aprobar haciendo lo '
    'fácil.',
    'Puntúa los 50. Un punto sin puntuar cuenta como cero en la media '
    'ponderada, y eso es a propósito: «no lo he mirado» no puede salir mejor '
    'que «lo he mirado y está mal».',
    'El área de prevención de riesgos laborales lleva su base legal punto por '
    'punto: temperatura del local, suelos, equipos de protección, resguardos '
    'de las máquinas, comprobaciones documentadas y formación del puesto. Cada '
    'uno con su norma y su enlace, comprobados el 6 de septiembre de 2026.',
    'Ojo con la temperatura del local: la norma da 17-27 °C para trabajo '
    'sedentario y 14-25 °C para trabajo ligero, y NO fija rango para el '
    'trabajo pesado. La categoría de trabajo de tu cocina la decide tu '
    'evaluación de riesgos, no una tabla general. ' + verificado('CE-21'),
    'Y con las cargas: el Real Decreto 487/1997 no contiene ninguna cifra en '
    'kilos. Los 25 kg que todo el mundo repite salen de una guía técnica que '
    'no es vinculante. ' + verificado('CE-24'),
    'La auditoría no sustituye al APPCC ni a la auditoría de servicio del '
    '«Manual del Manager de Restaurante». Son tres cosas distintas y cada una '
    'tiene su libro.',
]


def hoja_instrucciones(wb):
    ws = wb.create_sheet('Instrucciones', 0)
    anchos(ws, {'A': 110.0})
    motor.val(ws, 'A1', TITULO)
    ws['A1'].font = Font(bold=True, size=16, color=ORO)
    ws.row_dimensions[1].height = 30
    motor.val(ws, 'A2', SUBTITULO)
    motor.val(ws, 'A3', 'Para qué sirve: puntuar la disciplina de tu cocina '
                        'siempre igual, ver si mejora y comparar unas cocinas '
                        'con otras.')
    ws['A3'].font = Font(italic=True, size=9)

    seccion(ws, 'A5', 'Instrucciones de uso')
    fila = 6
    for paso in PASOS:
        nota(ws, fila, paso, alto=48, wrap=True)
        fila += 1
    fila += 1
    motor.val(ws, 'A' + str(fila), motor.NOTA_VERDES)
    ws['A' + str(fila)].fill = PatternFill('solid', fgColor=motor.VERDE)
    fila += 2
    seccion(ws, 'A' + str(fila), 'Lo que conviene saber antes de empezar')
    fila += 1
    for texto in NOTAS_LIBRO:
        nota(ws, fila, texto, alto=62, wrap=True)
        fila += 1
    fila += 1
    motor.val(ws, 'A' + str(fila), D.NOTA_DESPROTEGER)
    fila += 1
    motor.val(ws, 'A' + str(fila),
              'Ruta en Excel: Revisar > Desproteger hoja. La protección no '
              'tiene contraseña; sirve para no pisar una fórmula sin querer.')
    fila += 2
    motor.val(ws, 'A' + str(fila), D.BIO)
    fila += 1
    motor.val(ws, 'A' + str(fila), D.VERSION_LINE)
    pagina(ws, apaisado=False)
    return ws


# --------------------------------------------------------------------------
COLS_AUD = [
    ('A', 'Nº', 'ent'),
    ('B', 'Área', 'txt'),
    ('C', 'Punto de control', 'txt'),
    ('D', 'Peso (1 a 3)', 'ent'),
    ('E', 'Puntuación (0 a la escala)', 'ent'),
    ('F', 'Ponderado', 'dec'),
    ('G', 'Norma aplicable', 'txt'),
    ('H', 'Verificación', 'txt'),
    ('I', 'Observación', 'txt'),
]


def hoja_auditoria(wb):
    ws = wb.create_sheet('Auditoría')
    encabezar(ws, 'Auditoría interna de cocina: 50 puntos en cinco áreas',
              nota='La nota es una media PONDERADA por el peso de cada punto. '
                   'No mira limpieza, plagas ni temperaturas de cámara: eso es '
                   'el Pack APPCC.')
    anchos(ws, {'A': 6, 'B': 26, 'C': 74, 'D': 11, 'E': 14, 'F': 11,
                'G': 46, 'H': 76, 'I': 40})

    seccion(ws, 'A5', 'ESTA VISITA')
    cabecera(ws, 6, [('B', 'Campo'), ('D', 'Valor')], altura=20)
    pintar_cabecera(ws, 6, 'CEF')
    ws.merge_cells('B6:C6')
    ws.merge_cells('D6:F6')

    visita = D.AUDITORIAS_COCINA_HECHAS[-1]
    entradas = [
        (A_FECHA, 'Fecha de la visita', D._fecha(visita[1]), FMT_FECHA),
        (A_AUDITOR, 'Quién audita', visita[2], None),
        (A_UNIDAD, 'Unidad o cocina auditada', D.UNIDADES[0][0], None),
        (A_ESCALA, 'Escala máxima de puntuación', 5.0, FMT_ENT),
        (A_AMBAR, 'Aviso ÁMBAR: puntuación por debajo de', 3.0, FMT_ENT),
        (A_ROJO, 'Aviso CRÍTICO: puntuación por debajo de', 2.0, FMT_ENT),
    ]
    for fila, texto, valor, fmt in entradas:
        motor.val(ws, 'B%d' % fila, texto)
        ws.merge_cells('B%d:C%d' % (fila, fila))
        motor.val(ws, 'D%d' % fila, valor, fmt=fmt, verde_=True)
        ws.merge_cells('D%d:F%d' % (fila, fila))
    motor.dv_fecha(ws, ['D%d' % A_FECHA])
    motor.dv_numerica(ws, ['D%d' % A_ESCALA], minimo=1, maximo=10,
                      titulo='Escala máxima')
    motor.dv_numerica(ws, ['D%d' % A_AMBAR, 'D%d' % A_ROJO], minimo=0,
                      maximo=10, titulo='Umbral del semáforo')

    calculadas = [
        (A_MEDIA, 'Puntuación media PONDERADA de la visita',
         '=IFERROR(SUMPRODUCT($D${a}:$D${b},$E${a}:$E${b})/SUM($D${a}:$D${b}),'
         '"")'.format(a=A_INI, b=A_FIN), FMT_DEC),
        (A_CUMP, 'Cumplimiento sobre la escala',
         '=IFERROR(IF(OR(NOT(ISNUMBER($D${m})),NOT(ISNUMBER($D${e})),'
         '$D${e}=0),"",$D${m}/$D${e}),"")'.format(m=A_MEDIA, e=A_ESCALA),
         FMT_PCT),
        (A_PUNT, 'Puntos puntuados (de los 50)',
         '=COUNT($E${a}:$E${b})'.format(a=A_INI, b=A_FIN), FMT_ENT),
    ]
    for fila, texto, formula, fmt in calculadas:
        motor.val(ws, 'B%d' % fila, texto, bold=True)
        ws.merge_cells('B%d:C%d' % (fila, fila))
        motor.f(ws, 'D%d' % fila, formula, fmt=fmt, bold=True)
        ws.merge_cells('D%d:F%d' % (fila, fila))
        gris(ws, 'D%d' % fila)
    nota(ws, 16,
         'La media ponderada divide la suma de peso x puntuación entre la suma '
         'de los pesos. Un punto sin puntuar cuenta como cero: puntúa los 50.',
         col='B')

    cabecera(ws, A_CAB, [(c, t) for c, t, _ in COLS_AUD], altura=44)
    punt = list(visita[3])
    for i, fila in enumerate(D.AUDITORIA_COCINA):
        r = A_INI + i
        num, area, texto, peso, idd, norma, url = fila
        motor.val(ws, 'A%d' % r, num, fmt=FMT_ENT, align='center')
        motor.val(ws, 'B%d' % r, area)
        motor.val(ws, 'C%d' % r, texto, wrap=True)
        motor.val(ws, 'D%d' % r, float(peso), fmt=FMT_ENT, verde_=True,
                  align='center')
        motor.val(ws, 'E%d' % r, float(punt[i]), fmt=FMT_ENT, verde_=True,
                  align='center')
        if idd:
            motor.val(ws, 'G%d' % r, norma, wrap=True)
            motor.val(ws, 'H%d' % r, verificado(idd), wrap=True)
        motor.verde(ws, 'I%d' % r)
        ws.row_dimensions[r].height = 30
        motor.f(ws, 'F%d' % r,
                '=IFERROR(IF(OR(NOT(ISNUMBER($D{r})),NOT(ISNUMBER($E{r}))),'
                '"",$D{r}*$E{r}),"")'.format(r=r), fmt=FMT_DEC, align='center')

    filas = list(range(A_INI, A_FIN + 1))
    motor.dv_numerica(ws, ['D%d' % r for r in filas], minimo=1, maximo=3,
                      titulo='Peso', mensaje='El peso va de 1 a 3.')
    motor.dv_numerica(ws, ['E%d' % r for r in filas], minimo=0, maximo=10,
                      titulo='Puntuación',
                      mensaje='Puntúa de 0 a la escala máxima que hayas '
                              'puesto arriba.')
    rango = 'E{a}:E{b}'.format(a=A_INI, b=A_FIN)
    cf_expresion(ws, rango,
                 '=AND(ISNUMBER($E{a}),ISNUMBER($D${r}),$E{a}<$D${r})'
                 .format(a=A_INI, r=A_ROJO), motor.CF_ROJO_BG,
                 motor.CF_ROJO_FG)
    cf_expresion(ws, rango,
                 '=AND(ISNUMBER($E{a}),ISNUMBER($D${m}),$E{a}<$D${m})'
                 .format(a=A_INI, m=A_AMBAR), motor.CF_AMBAR_BG,
                 motor.CF_AMBAR_FG)
    cf_expresion(ws, rango,
                 '=AND(ISNUMBER($E{a}),ISNUMBER($D${m}),$E{a}>=$D${m})'
                 .format(a=A_INI, m=A_AMBAR), motor.CF_VERDE_BG,
                 motor.CF_VERDE_FG)

    ws.freeze_panes = 'D%d' % A_INI
    pagina(ws, titulos='$%d:$%d' % (A_CAB, A_CAB))
    return ws


# --------------------------------------------------------------------------
COLS_RES = [
    ('A', 'Área', 'txt'),
    ('B', 'Puntos puntuados', 'ent'),
    ('C', 'Peso total del área', 'ent'),
    ('D', 'Nota de esta visita', 'dec'),
    ('E', 'Cumplimiento sobre la escala', 'pct'),
    ('F', 'Tu objetivo (0 a la escala)', 'dec'),
    ('G', 'Nota de la visita comparada', 'dec'),
    ('H', 'Variación', 'dec'),
    ('I', 'Lectura', 'txt'),
]


def hoja_resumen(wb):
    ws = wb.create_sheet('Resumen por Área')
    encabezar(ws, 'Resumen por área',
              nota='Una media global esconde el área que empeora. Aquí se ve '
                   'por separado, y comparada con la visita que elijas.')
    anchos(ws, {'A': 34, 'B': 12, 'C': 13, 'D': 15, 'E': 16, 'F': 16,
                'G': 18, 'H': 12, 'I': 40})

    motor.val(ws, 'A%d' % S_COMP,
              'Visita del histórico con la que comparar (nº)', bold=True)
    motor.val(ws, 'B%d' % S_COMP, 1.0, fmt=FMT_ENT, verde_=True,
              align='center')
    motor.dv_numerica(ws, ['B%d' % S_COMP], minimo=1, maximo=99,
                      titulo='Número de visita')
    motor.val(ws, 'C%d' % S_COMP,
              'Escribe aquí el número de una visita de la hoja «Histórico» y '
              'las columnas de comparación se recalculan.')

    cabecera(ws, S_CAB, [(c, t) for c, t, _ in COLS_RES], altura=44)
    for i, area in enumerate(AREAS):
        r = S_INI + i
        ini, fin = BLOQUES[i]
        a, b = A_INI + ini, A_INI + fin
        motor.val(ws, 'A%d' % r, area, bold=True)
        motor.f(ws, 'B%d' % r,
                '=COUNT({u}!$E${a}:$E${b})'.format(u=AUD, a=a, b=b),
                fmt=FMT_ENT, align='center')
        motor.f(ws, 'C%d' % r,
                '=SUM({u}!$D${a}:$D${b})'.format(u=AUD, a=a, b=b),
                fmt=FMT_ENT, align='center')
        motor.f(ws, 'D%d' % r,
                '=IFERROR(SUMPRODUCT({u}!$D${a}:$D${b},{u}!$E${a}:$E${b})'
                '/SUM({u}!$D${a}:$D${b}),"")'.format(u=AUD, a=a, b=b),
                fmt=FMT_DEC, bold=True)
        motor.f(ws, 'E%d' % r,
                '=IFERROR(IF(OR(NOT(ISNUMBER($D{r})),'
                'NOT(ISNUMBER({u}!$D${e})),{u}!$D${e}=0),"",'
                '$D{r}/{u}!$D${e}),"")'.format(r=r, u=AUD, e=A_ESCALA),
                fmt=FMT_PCT)
        motor.verde(ws, 'F%d' % r)
        ws['F%d' % r].number_format = FMT_DEC
        motor.f(ws, 'G%d' % r,
                '=IFERROR(INDEX({h}!$D${a}:$H${b},MATCH($B${c},'
                '{h}!$A${a}:$A${b},0),MATCH($A{r},{h}!$D${k}:$H${k},0)),"")'
                .format(h=HIS, a=H_INI, b=H_FIN, c=S_COMP, r=r, k=H_CAB),
                fmt=FMT_DEC)
        motor.f(ws, 'H%d' % r,
                '=IFERROR(IF(OR(NOT(ISNUMBER($D{r})),NOT(ISNUMBER($G{r}))),'
                '"",$D{r}-$G{r}),"")'.format(r=r), fmt=FMT_DEC, bold=True)
        motor.f(ws, 'I%d' % r,
                '=IF(NOT(ISNUMBER($H{r})),"",IF($H{r}<0,"{mal}","{bien}"))'
                .format(r=r, mal=LECT_MAL, bien=LECT_BIEN))
        ws.row_dimensions[r].height = 26
    motor.dv_numerica(ws, ['F%d' % r for r in range(S_INI, S_FIN + 1)],
                      minimo=0, maximo=10, titulo='Tu objetivo')

    motor.val(ws, 'A%d' % S_TOT, 'TOTAL PONDERADO DE LA VISITA', bold=True)
    motor.f(ws, 'B%d' % S_TOT,
            '=COUNT({u}!$E${a}:$E${b})'.format(u=AUD, a=A_INI, b=A_FIN),
            fmt=FMT_ENT, align='center')
    motor.f(ws, 'C%d' % S_TOT,
            '=SUM({u}!$D${a}:$D${b})'.format(u=AUD, a=A_INI, b=A_FIN),
            fmt=FMT_ENT, align='center')
    motor.f(ws, 'D%d' % S_TOT,
            '=IF({u}!$D${m}="","",{u}!$D${m})'.format(u=AUD, m=A_MEDIA),
            fmt=FMT_DEC, bold=True)
    motor.f(ws, 'E%d' % S_TOT,
            '=IF({u}!$D${c}="","",{u}!$D${c})'.format(u=AUD, c=A_CUMP),
            fmt=FMT_PCT)
    motor.f(ws, 'G%d' % S_TOT,
            '=IFERROR(INDEX({h}!$I${a}:$I${b},MATCH($B${c},{h}!$A${a}:$A${b},'
            '0)),"")'.format(h=HIS, a=H_INI, b=H_FIN, c=S_COMP), fmt=FMT_DEC)
    motor.f(ws, 'H%d' % S_TOT,
            '=IFERROR(IF(OR(NOT(ISNUMBER($D${t})),NOT(ISNUMBER($G${t}))),"",'
            '$D${t}-$G${t}),"")'.format(t=S_TOT), fmt=FMT_DEC, bold=True)
    motor.f(ws, 'I%d' % S_TOT,
            '=IF(NOT(ISNUMBER($H${t})),"",IF($H${t}<0,"{mal}","{bien}"))'
            .format(t=S_TOT, mal=LECT_MAL, bien=LECT_BIEN))
    fila_total(ws, S_TOT, 'A', 'I')

    rango = 'H{a}:H{b}'.format(a=S_INI, b=S_TOT)
    cf_expresion(ws, rango,
                 '=AND(ISNUMBER($H{a}),$H{a}<0)'.format(a=S_INI),
                 motor.CF_ROJO_BG, motor.CF_ROJO_FG)
    cf_expresion(ws, rango,
                 '=AND(ISNUMBER($H{a}),$H{a}>=0)'.format(a=S_INI),
                 motor.CF_VERDE_BG, motor.CF_VERDE_FG)
    motor.semaforo_texto(ws, 'I{a}:I{b}'.format(a=S_INI, b=S_TOT),
                         ((LECT_MAL, motor.CF_ROJO_BG, motor.CF_ROJO_FG),
                          (LECT_BIEN, motor.CF_VERDE_BG,
                           motor.CF_VERDE_FG)))
    for r in range(S_INI, S_FIN + 1):
        cf_expresion(ws, 'D%d' % r,
                     '=AND(ISNUMBER($D{r}),ISNUMBER($F{r}),$D{r}<$F{r})'
                     .format(r=r), motor.CF_ROJO_BG, motor.CF_ROJO_FG)
        cf_expresion(ws, 'D%d' % r,
                     '=AND(ISNUMBER($D{r}),ISNUMBER($F{r}),$D{r}>=$F{r})'
                     .format(r=r), motor.CF_VERDE_BG, motor.CF_VERDE_FG)

    fila = S_TOT + 2
    for texto in (
            'La columna «Tu objetivo» viene VACÍA a propósito. Mientras la '
            'dejes vacía su semáforo no se enciende: no existe ninguna nota de '
            'auditoría de cocina publicada en español con la que compararte, y '
            'aquí no se inventa una.',
            'La comparación que sí vale es contra tu visita anterior, y por '
            'ÁREA. En el juego de ejemplo la nota global sube y el área de '
            'mermas por partida cae más de un punto: eso es exactamente lo que '
            'una media global esconde.',
            'El peso de cada punto se lee de la hoja «Auditoría»: si lo '
            'cambias allí, estas notas se recalculan solas.'):
        nota(ws, fila, texto, alto=42, wrap=True)
        ws.merge_cells('A%d:I%d' % (fila, fila))
        fila += 1

    pagina(ws)
    return ws


# --------------------------------------------------------------------------
def hoja_historico(wb):
    ws = wb.create_sheet('Histórico')
    encabezar(ws, 'Histórico de auditorías de esta cocina',
              nota='Una fila por visita. La visita en curso se lee sola de las '
                   'otras hojas; las cerradas se pegan como valores.')
    anchos(ws, {'A': 12, 'B': 14, 'C': 34, 'D': 17, 'E': 17, 'F': 17,
                'G': 17, 'H': 17, 'I': 17, 'J': 16, 'K': 30, 'L': 18})
    letras = ['D', 'E', 'F', 'G', 'H']

    motor.val(ws, 'A%d' % H_PESOS,
              'Peso total de cada área (se lee del resumen)', bold=True)
    ws.merge_cells('A%d:C%d' % (H_PESOS, H_PESOS))
    for i, letra in enumerate(letras):
        motor.f(ws, '%s%d' % (letra, H_PESOS),
                '=IF({r}!$C${f}="","",{r}!$C${f})'.format(r=RES, f=S_INI + i),
                fmt=FMT_ENT, align='center')
        gris(ws, '%s%d' % (letra, H_PESOS))

    cabecera(ws, H_CAB,
             [('A', 'Nº de visita'), ('B', 'Fecha'), ('C', 'Quién audita')]
             + [(letras[i], AREAS[i]) for i in range(len(AREAS))]
             + [('I', 'Nota ponderada global'),
                ('J', 'Cumplimiento sobre la escala'),
                ('K', 'Unidad o cocina'),
                ('L', 'Variación respecto a la fila anterior')], altura=58)

    # --- fila 1: la visita ya cerrada, con sus notas por área --------------
    v1 = D.AUDITORIAS_COCINA_HECHAS[0]
    motor.val(ws, 'A%d' % H_INI, float(v1[0]), fmt=FMT_ENT, verde_=True,
              align='center')
    motor.val(ws, 'B%d' % H_INI, D._fecha(v1[1]), fmt=FMT_FECHA, verde_=True)
    motor.val(ws, 'C%d' % H_INI, v1[2], verde_=True)
    pesos = [f[3] for f in D.AUDITORIA_COCINA]
    for i, letra in enumerate(letras):
        ini, fin = BLOQUES[i]
        num = sum(v1[3][j] * pesos[j] for j in range(ini, fin + 1))
        den = sum(pesos[j] for j in range(ini, fin + 1))
        # seis decimales, no cuatro: con cuatro, la nota global que este libro
        # recompone de la visita cerrada se separaba 5 millonésimas de la que
        # sale de sus 50 puntuaciones, y un refutador lo lee como un descuadre
        motor.val(ws, '%s%d' % (letra, H_INI), round(num / float(den), 6),
                  fmt=FMT_DEC, verde_=True)
    motor.val(ws, 'K%d' % H_INI, D.UNIDADES[0][0], verde_=True)

    # --- fila 2: la visita EN CURSO, enlazada -----------------------------
    r2 = H_INI + 1
    motor.val(ws, 'A%d' % r2, float(D.AUDITORIAS_COCINA_HECHAS[-1][0]),
              fmt=FMT_ENT, verde_=True, align='center')
    motor.f(ws, 'B%d' % r2,
            '=IF({u}!$D${f}="","",{u}!$D${f})'.format(u=AUD, f=A_FECHA),
            fmt=FMT_FECHA)
    motor.f(ws, 'C%d' % r2,
            '=IF({u}!$D${f}="","",{u}!$D${f})'.format(u=AUD, f=A_AUDITOR))
    for i, letra in enumerate(letras):
        motor.f(ws, '%s%d' % (letra, r2),
                '=IF({r}!$D${f}="","",{r}!$D${f})'.format(r=RES, f=S_INI + i),
                fmt=FMT_DEC)
    motor.f(ws, 'K%d' % r2,
            '=IF({u}!$D${f}="","",{u}!$D${f})'.format(u=AUD, f=A_UNIDAD))
    for letra in ('B', 'C', 'K'):
        gris(ws, '%s%d' % (letra, r2))
    motor.val(ws, 'M%d' % r2, 'VISITA EN CURSO: esta fila se lee sola de las '
                              'hojas «Auditoría» y «Resumen por Área».')
    ws.column_dimensions['M'].width = 20

    for r in range(H_INI, H_FIN + 1):
        if r > r2:
            motor.verde(ws, 'A%d:H%d' % (r, r))
            motor.verde(ws, 'K%d' % r)
        motor.f(ws, 'I%d' % r,
                '=IFERROR(IF(OR(COUNT($D{r}:$H{r})=0,SUM($D${p}:$H${p})=0),'
                '"",SUMPRODUCT($D{r}:$H{r},$D${p}:$H${p})'
                '/SUM($D${p}:$H${p})),"")'.format(r=r, p=H_PESOS),
                fmt=FMT_DEC, bold=True)
        motor.f(ws, 'J%d' % r,
                '=IFERROR(IF(OR(NOT(ISNUMBER($I{r})),'
                'NOT(ISNUMBER({u}!$D${e})),{u}!$D${e}=0),"",'
                '$I{r}/{u}!$D${e}),"")'.format(r=r, u=AUD, e=A_ESCALA),
                fmt=FMT_PCT)
        if r > H_INI:
            motor.f(ws, 'L%d' % r,
                    '=IFERROR(IF(OR(NOT(ISNUMBER($I{r})),'
                    'NOT(ISNUMBER($I{p}))),"",$I{r}-$I{p}),"")'
                    .format(r=r, p=r - 1), fmt=FMT_DEC, bold=True)
    motor.dv_fecha(ws, ['B%d' % r for r in range(r2 + 1, H_FIN + 1)])
    motor.dv_numerica(ws, ['%s%d' % (c, r) for c in letras
                           for r in range(r2 + 1, H_FIN + 1)],
                      minimo=0, maximo=10, titulo='Nota del área')
    motor.dv_numerica(ws, ['A%d' % r for r in range(H_INI, H_FIN + 1)],
                      minimo=1, maximo=99, titulo='Número de visita')

    rango = 'L{a}:L{b}'.format(a=H_INI, b=H_FIN)
    cf_expresion(ws, rango, '=AND(ISNUMBER($L{a}),$L{a}<0)'.format(a=H_INI),
                 motor.CF_ROJO_BG, motor.CF_ROJO_FG)
    cf_expresion(ws, rango, '=AND(ISNUMBER($L{a}),$L{a}>=0)'.format(a=H_INI),
                 motor.CF_VERDE_BG, motor.CF_VERDE_FG)

    fila = H_FIN + 2
    for texto in (
            'Cómo se cierra una visita: cuando termines la auditoría, copia la '
            'fila de la visita en curso y pégala como VALORES en la primera '
            'fila libre de abajo. A partir de ahí la fila enlazada vuelve a '
            'quedar libre para la siguiente.',
            'La nota global de cada fila NO es la media de las cinco áreas: es '
            'la media ponderada por el peso de cada área, que está en la fila '
            'de arriba. Por eso un área con mucho peso mueve más la nota.',
            'Guarda también quién audita. Una autoauditoría del jefe de cocina '
            'y una visita de dirección con el chef delante no puntúan igual, y '
            'saber cuál fue cada una evita conclusiones falsas.'):
        nota(ws, fila, texto, alto=42, wrap=True)
        ws.merge_cells('A%d:L%d' % (fila, fila))
        fila += 1

    ws.freeze_panes = 'D%d' % H_INI
    pagina(ws, titulos='$%d:$%d' % (H_CAB, H_CAB))
    return ws


# --------------------------------------------------------------------------
def hoja_unidades(wb):
    ws = wb.create_sheet('Histórico por Unidad')
    encabezar(ws, 'La misma auditoría, cocina por cocina',
              nota='«Una cocina, con las herramientas para comparar varias». '
                   'La unidad de referencia se calcula sola; las demás las '
                   'rellenas cuando las audites.')
    anchos(ws, {'A': 34, 'B': 18, 'C': 17, 'D': 17, 'E': 17, 'F': 17,
                'G': 17, 'H': 17, 'I': 16, 'J': 20, 'K': 20, 'L': 34})
    letras = ['C', 'D', 'E', 'F', 'G']

    motor.val(ws, 'A%d' % U_PESOS,
              'Peso total de cada área (se lee del resumen)', bold=True)
    ws.merge_cells('A%d:B%d' % (U_PESOS, U_PESOS))
    for i, letra in enumerate(letras):
        motor.f(ws, '%s%d' % (letra, U_PESOS),
                '=IF({r}!$C${f}="","",{r}!$C${f})'.format(r=RES, f=S_INI + i),
                fmt=FMT_ENT, align='center')
        gris(ws, '%s%d' % (letra, U_PESOS))

    cabecera(ws, U_CAB,
             [('A', 'Unidad o cocina'), ('B', 'Fecha de la auditoría')]
             + [(letras[i], AREAS[i]) for i in range(len(AREAS))]
             + [('H', 'Nota ponderada'),
                ('I', 'Cumplimiento sobre la escala'),
                ('J', 'Desviación contra la media del grupo'),
                ('K', 'Desviación contra tu estándar'),
                ('L', 'Lectura')], altura=58)

    # fila 1: la unidad de referencia, enlazada al resumen. A2 de la
    # refutación xlsx (2026-09-06): antes era un literal que no se enteraba
    # de que cambiaras el nombre de tu cocina en «Auditoría».
    motor.f(ws, 'A%d' % U_INI,
            '=IF({u}!$D${f}="","",{u}!$D${f})'.format(u=AUD, f=A_UNIDAD))
    gris(ws, 'A%d' % U_INI)
    motor.f(ws, 'B%d' % U_INI,
            '=IF({u}!$D${f}="","",{u}!$D${f})'.format(u=AUD, f=A_FECHA),
            fmt=FMT_FECHA)
    gris(ws, 'B%d' % U_INI)
    for i, letra in enumerate(letras):
        motor.f(ws, '%s%d' % (letra, U_INI),
                '=IF({r}!$D${f}="","",{r}!$D${f})'.format(r=RES, f=S_INI + i),
                fmt=FMT_DEC)
    # las demás unidades del grupo: nombre sí, puntuación NO (no se inventa)
    for i, unidad in enumerate(D.UNIDADES[1:]):
        r = U_INI + 1 + i
        motor.val(ws, 'A%d' % r, unidad[0], verde_=True)

    for r in range(U_INI, U_FIN + 1):
        if r > U_INI:
            motor.verde(ws, 'A%d:G%d' % (r, r))
        motor.f(ws, 'H%d' % r,
                '=IFERROR(IF(OR(COUNT($C{r}:$G{r})=0,SUM($C${p}:$G${p})=0),'
                '"",SUMPRODUCT($C{r}:$G{r},$C${p}:$G${p})'
                '/SUM($C${p}:$G${p})),"")'.format(r=r, p=U_PESOS),
                fmt=FMT_DEC, bold=True)
        motor.f(ws, 'I%d' % r,
                '=IFERROR(IF(OR(NOT(ISNUMBER($H{r})),'
                'NOT(ISNUMBER({u}!$D${e})),{u}!$D${e}=0),"",'
                '$H{r}/{u}!$D${e}),"")'.format(r=r, u=AUD, e=A_ESCALA),
                fmt=FMT_PCT)
        motor.f(ws, 'J%d' % r,
                '=IFERROR(IF(OR(NOT(ISNUMBER($H{r})),NOT(ISNUMBER($C${m}))),'
                '"",$H{r}-$C${m}),"")'.format(r=r, m=U_MEDIA), fmt=FMT_DEC)
        motor.f(ws, 'K%d' % r,
                '=IFERROR(IF(OR(NOT(ISNUMBER($H{r})),NOT(ISNUMBER($C${e}))),'
                '"",$H{r}-$C${e}),"")'.format(r=r, e=U_ESTANDAR), fmt=FMT_DEC)
        motor.f(ws, 'L%d' % r,
                '=IF(NOT(ISNUMBER($K{r})),"",IF($K{r}<0,"{mal}","{bien}"))'
                .format(r=r, mal=LECT_OBJ_MAL, bien=LECT_OBJ_BIEN))
        ws.row_dimensions[r].height = 24
    motor.dv_fecha(ws, ['B%d' % r for r in range(U_INI + 1, U_FIN + 1)])
    motor.dv_numerica(ws, ['%s%d' % (c, r) for c in letras
                           for r in range(U_INI + 1, U_FIN + 1)],
                      minimo=0, maximo=10, titulo='Nota del área')

    motor.val(ws, 'A%d' % U_MEDIA, 'Media del grupo (unidades auditadas)',
              bold=True)
    ws.merge_cells('A%d:B%d' % (U_MEDIA, U_MEDIA))
    motor.f(ws, 'C%d' % U_MEDIA,
            '=IFERROR(IF(COUNT($H${a}:$H${b})=0,"",SUM($H${a}:$H${b})'
            '/COUNT($H${a}:$H${b})),"")'.format(a=U_INI, b=U_FIN),
            fmt=FMT_DEC, bold=True)
    gris(ws, 'C%d' % U_MEDIA)
    motor.val(ws, 'A%d' % U_ESTANDAR,
              'Tu estándar de grupo (objetivo propio, 0 a la escala)',
              bold=True)
    ws.merge_cells('A%d:B%d' % (U_ESTANDAR, U_ESTANDAR))
    motor.val(ws, 'C%d' % U_ESTANDAR, ESTANDAR_GRUPO_EJEMPLO, fmt=FMT_DEC,
              verde_=True)
    motor.dv_numerica(ws, ['C%d' % U_ESTANDAR], minimo=0, maximo=10,
                      titulo='Tu estándar')

    rango = 'J{a}:J{b}'.format(a=U_INI, b=U_FIN)
    cf_expresion(ws, rango, '=AND(ISNUMBER($J{a}),$J{a}<0)'.format(a=U_INI),
                 motor.CF_ROJO_BG, motor.CF_ROJO_FG)
    cf_expresion(ws, rango, '=AND(ISNUMBER($J{a}),$J{a}>=0)'.format(a=U_INI),
                 motor.CF_VERDE_BG, motor.CF_VERDE_FG)
    motor.semaforo_texto(ws, 'L{a}:L{b}'.format(a=U_INI, b=U_FIN),
                         ((LECT_OBJ_MAL, motor.CF_ROJO_BG, motor.CF_ROJO_FG),
                          (LECT_OBJ_BIEN, motor.CF_VERDE_BG,
                           motor.CF_VERDE_FG)))

    fila = U_ESTANDAR + 2
    for texto in (
            'Las unidades del grupo vienen con NOMBRE y sin puntuación: sus '
            'notas son las que saques tú cuando las audites. Aquí no se '
            'inventa la puntuación de una cocina que nadie ha visitado.',
            'Para que la comparación valga, las tres cosas tienen que ser '
            'iguales: los mismos 50 puntos, los mismos pesos y la misma '
            'escala. Si en una unidad cambias un peso, cámbialo en todas o '
            'deja de comparar.',
            'Comparar cocinas de formatos distintos sin decirlo es la forma '
            'más rápida de tomar una decisión mala con un número bueno. Una '
            'cocina de banquetes y una de carta no se auditan igual aunque el '
            'checklist sea el mismo.',
            'El ejemplo trae el estándar sembrado con la nota ponderada de tu '
            'última visita (%.2f): empezar exigiéndote no bajar de donde ya '
            'estás es un criterio razonable. Cámbialo por el tuyo cuando '
            'quieras.' % ESTANDAR_GRUPO_EJEMPLO):
        nota(ws, fila, texto, alto=42, wrap=True)
        ws.merge_cells('A%d:L%d' % (fila, fila))
        fila += 1

    pagina(ws, titulos='$%d:$%d' % (U_CAB, U_CAB))
    return ws


# --------------------------------------------------------------------------
COLS_NORM = [
    ('A', 'Norma', 'txt'),
    ('B', 'Estado a la fecha de corte', 'txt'),
    ('C', 'Qué tiene que hacer el chef', 'txt'),
    ('D', 'Fecha de corte', 'fecha'),
    ('E', 'Id', 'txt'),
    ('F', 'Verificación y enlace', 'txt'),
    ('G', 'Comprobado por ti el', 'fecha'),
    ('H', 'Estado tras tu comprobación', 'txt'),
]


def hoja_normativo(wb):
    ws = wb.create_sheet('Estado Normativo')
    encabezar(ws, 'Estado normativo a la fecha de corte',
              nota='Qué dice hoy cada norma que toca esta auditoría, con el '
                   'enlace para que lo compruebes tú. La fecha de corte es '
                   'editable.')
    anchos(ws, {'A': 46, 'B': 40, 'C': 64, 'D': 15, 'E': 9, 'F': 76,
                'G': 17, 'H': 24})

    motor.val(ws, 'A%d' % N_CORTE, 'Fecha de corte de esta revisión',
              bold=True)
    motor.val(ws, 'B%d' % N_CORTE,
              D._fecha(D.PARAMETROS_COCINA['fecha_corte_normativa']),
              fmt=FMT_FECHA, verde_=True)
    motor.dv_fecha(ws, ['B%d' % N_CORTE])
    motor.val(ws, 'C%d' % N_CORTE,
              'Es el día en que se comprobó cada norma contra el texto '
              'consolidado del BOE o de EUR-Lex.')

    motor.val(ws, 'A%d' % N_SIM,
              'Fecha para simular (déjala vacía y se usa la de hoy)')
    motor.val(ws, 'B%d' % N_SIM, None, fmt=FMT_FECHA, verde_=True)
    motor.dv_fecha(ws, ['B%d' % N_SIM])

    motor.val(ws, 'A%d' % N_HOY, 'Hoy (fecha de referencia)', bold=True)
    motor.f(ws, 'B%d' % N_HOY,
            '=IFERROR(IF($B${s}="",TODAY(),$B${s}),"")'.format(s=N_SIM),
            fmt=FMT_FECHA, bold=True)
    gris(ws, 'B%d' % N_HOY)

    motor.val(ws, 'A%d' % N_DIAS, 'Días naturales desde la fecha de corte')
    motor.f(ws, 'B%d' % N_DIAS,
            '=IFERROR(IF(OR(NOT(ISNUMBER($B${h})),NOT(ISNUMBER($B${c}))),"",'
            '$B${h}-$B${c}),"")'.format(h=N_HOY, c=N_CORTE), fmt=FMT_ENT)
    gris(ws, 'B%d' % N_DIAS)

    motor.val(ws, 'A%d' % N_UMBRAL,
              'Avisar cuando hayan pasado más de (días)')
    motor.val(ws, 'B%d' % N_UMBRAL, float(DIAS_AVISO_REVISION), fmt=FMT_ENT,
              verde_=True)
    motor.dv_numerica(ws, ['B%d' % N_UMBRAL], minimo=1, maximo=3650,
                      titulo='Días de aviso')
    motor.val(ws, 'C%d' % N_UMBRAL,
              'Criterio de la casa, no una obligación de ninguna norma: media '
              'año es un plazo razonable para volver a mirar los enlaces.')

    motor.val(ws, 'A%d' % N_AVISO, 'Aviso', bold=True)
    motor.f(ws, 'B%d' % N_AVISO,
            '=IF(OR(NOT(ISNUMBER($B${d})),NOT(ISNUMBER($B${u}))),"",'
            'IF($B${d}>$B${u},"{viejo}","{nuevo}"))'
            .format(d=N_DIAS, u=N_UMBRAL, viejo=AVISO_VIEJO,
                    nuevo=AVISO_NUEVO), bold=True)
    ws.merge_cells('B%d:C%d' % (N_AVISO, N_AVISO))
    motor.semaforo_texto(ws, 'B%d' % N_AVISO,
                         ((AVISO_VIEJO, motor.CF_ROJO_BG, motor.CF_ROJO_FG),
                          (AVISO_NUEVO, motor.CF_VERDE_BG,
                           motor.CF_VERDE_FG)))

    cabecera(ws, N_CAB, [(c, t) for c, t, _ in COLS_NORM], altura=44)
    for i, fila in enumerate(D.ESTADO_NORMATIVO):
        r = N_INI + i
        norma, estado, que_hace, _fecha, idd, _url = fila
        motor.val(ws, 'A%d' % r, norma, wrap=True)
        motor.val(ws, 'B%d' % r, estado, wrap=True)
        motor.val(ws, 'C%d' % r, que_hace, wrap=True)
        motor.val(ws, 'E%d' % r, idd, align='center')
        motor.val(ws, 'F%d' % r, verificado(idd), wrap=True)
        motor.verde(ws, 'G%d' % r)
        ws['G%d' % r].number_format = FMT_FECHA
        motor.val(ws, 'H%d' % r, ESTADOS_REVISION[0], verde_=True)
        ws.row_dimensions[r].height = 44
        motor.f(ws, 'D%d' % r,
                '=IF($B${c}="","",$B${c})'.format(c=N_CORTE), fmt=FMT_FECHA)

    filas = list(range(N_INI, N_FIN + 1))
    motor.dv_fecha(ws, ['G%d' % r for r in filas])
    dv_rango(ws, ['H%d' % r for r in filas], REF_ESTADOS_REVISION,
             'Estado tras tu comprobación',
             'Elige un estado de la lista del final de esta hoja.')
    motor.semaforo_texto(ws, 'H{a}:H{b}'.format(a=N_INI, b=N_FIN),
                         ((ESTADOS_REVISION[2], motor.CF_ROJO_BG,
                           motor.CF_ROJO_FG),
                          (ESTADOS_REVISION[1], motor.CF_VERDE_BG,
                           motor.CF_VERDE_FG),
                          (ESTADOS_REVISION[0], motor.CF_GRIS_BG,
                           motor.CF_GRIS_FG)))

    fila = N_FIN + 2
    for texto in (
            'Esta hoja existe para que no tengas que creerte nada. Cada fila '
            'trae el enlace al texto consolidado: si algo ha cambiado desde el '
            'día de corte, lo ves tú en dos clics y lo anotas en las dos '
            'últimas columnas.',
            'Las normas de esta hoja son las que toca la auditoría de cocina: '
            'seguridad alimentaria, alérgenos, comidas testigo, temperaturas, '
            'anisakis, aceites de fritura, prevención de riesgos y '
            'desperdicio. Las de convenio y las laborales viven en el '
            '«Manual del Manager de Restaurante».',
            'Que una norma esté vigente no quiere decir que te aplique entera. '
            'La columna «Qué tiene que hacer el chef» es la que traduce el '
            'artículo a la cocina; el resto es contexto.'):
        nota(ws, fila, texto, alto=44, wrap=True)
        ws.merge_cells('A%d:H%d' % (fila, fila))
        fila += 1

    # --- lista del desplegable (A1 de la refutación xlsx) ------------------
    seccion(ws, 'A%d' % N_LSEC, 'LISTA DEL DESPLEGABLE (no la borres)')
    cabecera(ws, N_REV_CAB, [('A', 'Estado tras tu comprobación')], altura=20)
    for i, texto in enumerate(ESTADOS_REVISION):
        motor.val(ws, 'A%d' % (N_REV_INI + i), texto)

    ws.freeze_panes = 'B%d' % N_INI
    pagina(ws, titulos='$%d:$%d' % (N_CAB, N_CAB))
    return ws


# --------------------------------------------------------------------------
def mapa():
    celdas_a = {
        'Fecha de la visita': 'D%d' % A_FECHA,
        'Quién audita': 'D%d' % A_AUDITOR,
        'Unidad o cocina auditada': 'D%d' % A_UNIDAD,
        'Escala máxima de puntuación': 'D%d' % A_ESCALA,
        'Umbral de aviso ÁMBAR': 'D%d' % A_AMBAR,
        'Umbral de aviso CRÍTICO': 'D%d' % A_ROJO,
        'Puntuación media PONDERADA de la visita': 'D%d' % A_MEDIA,
        'Cumplimiento sobre la escala': 'D%d' % A_CUMP,
        'Puntos puntuados': 'D%d' % A_PUNT,
    }
    for i, fila in enumerate(D.AUDITORIA_COCINA):
        celdas_a['Ponderado del punto %d' % fila[0]] = 'F%d' % (A_INI + i)

    celdas_s = {
        'Visita del histórico con la que comparar': 'B%d' % S_COMP,
        'Puntos puntuados en total': 'B%d' % S_TOT,
        'Peso total de los 50 puntos': 'C%d' % S_TOT,
        'Nota ponderada global de esta visita': 'D%d' % S_TOT,
        'Cumplimiento global sobre la escala': 'E%d' % S_TOT,
        'Nota global de la visita comparada': 'G%d' % S_TOT,
        'Variación global': 'H%d' % S_TOT,
        'Lectura global': 'I%d' % S_TOT,
    }
    for i, area in enumerate(AREAS):
        r = S_INI + i
        celdas_s['Puntos puntuados del área ' + area] = 'B%d' % r
        celdas_s['Peso total del área ' + area] = 'C%d' % r
        celdas_s['Nota de esta visita del área ' + area] = 'D%d' % r
        celdas_s['Cumplimiento del área ' + area] = 'E%d' % r
        celdas_s['Nota de la visita comparada del área ' + area] = 'G%d' % r
        celdas_s['Variación del área ' + area] = 'H%d' % r
        celdas_s['Lectura del área ' + area] = 'I%d' % r

    celdas_h = {}
    for i, area in enumerate(AREAS):
        celdas_h['Peso total del área ' + area] = '%s%d' % ('DEFGH'[i],
                                                            H_PESOS)
    for n, r in (('visita cerrada', H_INI), ('visita en curso', H_INI + 1)):
        celdas_h['Nota ponderada global de la ' + n] = 'I%d' % r
        celdas_h['Cumplimiento de la ' + n] = 'J%d' % r
    celdas_h['Variación de la visita en curso'] = 'L%d' % (H_INI + 1)
    for i, area in enumerate(AREAS):
        celdas_h['Nota de la visita en curso del área ' + area] = \
            '%s%d' % ('DEFGH'[i], H_INI + 1)

    celdas_u = {
        'Media del grupo': 'C%d' % U_MEDIA,
        'Tu estándar de grupo': 'C%d' % U_ESTANDAR,
    }
    for i, area in enumerate(AREAS):
        celdas_u['Peso total del área ' + area] = '%s%d' % ('CDEFG'[i],
                                                            U_PESOS)
        celdas_u['Nota de la unidad de referencia en el área ' + area] = \
            '%s%d' % ('CDEFG'[i], U_INI)
    celdas_u['Nota ponderada de la unidad de referencia'] = 'H%d' % U_INI
    celdas_u['Cumplimiento de la unidad de referencia'] = 'I%d' % U_INI
    celdas_u['Desviación de la unidad de referencia contra la media'] = \
        'J%d' % U_INI

    celdas_n = {
        'Fecha de corte de esta revisión': 'B%d' % N_CORTE,
        'Fecha para simular': 'B%d' % N_SIM,
        'Hoy (fecha de referencia)': 'B%d' % N_HOY,
        'Días naturales desde la fecha de corte': 'B%d' % N_DIAS,
        'Días tras los que avisar': 'B%d' % N_UMBRAL,
        'Aviso de revisión': 'B%d' % N_AVISO,
    }
    for i, fila in enumerate(D.ESTADO_NORMATIVO):
        celdas_n['Fecha de corte de ' + fila[4]] = 'D%d' % (N_INI + i)

    tipo = {'eur': 'eur', 'pct': 'pct1', 'ent': 'num', 'dec': 'num',
            'fecha': 'txt', 'txt': 'txt'}
    prl = [f for f in D.AUDITORIA_COCINA
           if f[4] and f[1] == AREAS[3]]
    return {
        'fichero': NOMBRE + '.xlsx',
        'producto': 'manual-chef-ejecutivo',
        'libro': 7,
        'que_decide': ('Puntuar la disciplina de cocina de forma repetible y '
                       'comparar cocinas'),
        'media_ponderada': {
            'formula': 'SUMPRODUCT(peso;puntuación)/SUM(peso)',
            'celda_global': 'Auditoría!D%d' % A_MEDIA,
            'celdas_por_area': dict((AREAS[i], 'Resumen por Área!D%d'
                                     % (S_INI + i)) for i in range(len(AREAS))),
            'puntos': len(D.AUDITORIA_COCINA),
            'peso_total': sum(f[3] for f in D.AUDITORIA_COCINA),
        },
        'prl_sembrada': [{'punto': f[0], 'id': f[4], 'norma': f[5],
                          'url': f[6],
                          'celda': 'Auditoría!C%d' % (A_INI + f[0] - 1)}
                         for f in prl],
        'estado_normativo': {
            'celda_fecha_de_corte': 'Estado Normativo!B%d' % N_CORTE,
            'editable': True,
            'normas': len(D.ESTADO_NORMATIVO),
        },
        'excluye': ('limpieza, control de plagas y temperaturas de cámara: '
                    'eso es el Pack APPCC'),
        'hojas': {
            'Auditoría': {
                'celdas': celdas_a,
                'tablas': [
                    {'titulo': 'Los 50 puntos de la auditoría',
                     'cols': [[t, c, tipo[k]] for c, t, k in COLS_AUD],
                     'filas': [A_INI, A_FIN]},
                ],
            },
            'Resumen por Área': {
                'celdas': celdas_s,
                'tablas': [
                    {'titulo': 'Resumen por área',
                     'cols': [[t, c, tipo[k]] for c, t, k in COLS_RES],
                     'filas': [S_INI, S_TOT]},
                ],
            },
            'Histórico': {
                'celdas': celdas_h,
                'tablas': [
                    {'titulo': 'Histórico de auditorías de esta cocina',
                     'cols': ([['Nº de visita', 'A', 'num'],
                               ['Fecha', 'B', 'txt'],
                               ['Quién audita', 'C', 'txt']]
                              + [[AREAS[i], 'DEFGH'[i], 'num']
                                 for i in range(len(AREAS))]
                              + [['Nota ponderada global', 'I', 'num'],
                                 ['Cumplimiento', 'J', 'pct1'],
                                 ['Unidad o cocina', 'K', 'txt'],
                                 ['Variación', 'L', 'num']]),
                     'filas': [H_INI, H_FIN]},
                ],
            },
            'Histórico por Unidad': {
                'celdas': celdas_u,
                'tablas': [
                    {'titulo': 'La misma auditoría, cocina por cocina',
                     'cols': ([['Unidad o cocina', 'A', 'txt'],
                               ['Fecha de la auditoría', 'B', 'txt']]
                              + [[AREAS[i], 'CDEFG'[i], 'num']
                                 for i in range(len(AREAS))]
                              + [['Nota ponderada', 'H', 'num'],
                                 ['Cumplimiento', 'I', 'pct1'],
                                 ['Desviación contra la media', 'J', 'num'],
                                 ['Desviación contra tu estándar', 'K', 'num'],
                                 ['Lectura', 'L', 'txt']]),
                     'filas': [U_INI, U_FIN]},
                ],
            },
            'Estado Normativo': {
                'celdas': celdas_n,
                'tablas': [
                    {'titulo': 'Estado normativo a la fecha de corte',
                     'cols': [[t, c, tipo[k]] for c, t, k in COLS_NORM],
                     'filas': [N_INI, N_FIN]},
                ],
            },
        },
    }


def main():
    wb = Workbook()
    wb.remove(wb.active)
    hoja_instrucciones(wb)
    hoja_auditoria(wb)
    hoja_resumen(wb)
    hoja_historico(wb)
    hoja_unidades(wb)
    hoja_normativo(wb)

    wb.properties.creator = 'AI Chef Pro'
    wb.properties.lastModifiedBy = 'AI Chef Pro'
    wb.properties.title = TITULO
    wb.properties.subject = PRODUCTO + ' · v1.0'

    for ws in wb.worksheets:
        motor.retirar_verde_de_calculadas(ws)
        motor.proteger(ws)

    destino = os.path.join(AQUI, 'build')
    os.makedirs(destino, exist_ok=True)
    ruta = os.path.join(destino, NOMBRE + '.xlsx')
    wb.save(ruta)
    with open(os.path.join(destino, 'mapa-' + NOMBRE + '.json'), 'w',
              encoding='utf-8') as fh:
        json.dump(mapa(), fh, ensure_ascii=False, indent=1)
    print('escrito:', ruta)
    print('formulas registradas:', len(motor.REGISTRO))


if __name__ == '__main__':
    main()
