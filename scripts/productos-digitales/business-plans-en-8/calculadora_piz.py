#!/usr/bin/env python3
"""
calculadora_piz.py — Pizza Dough & Food Cost Calculator (Neapolitan-Style Dough + Food Cost for 8 Pizzas), D145 del Pizza Shop
Business Plan Kit (EN, $39). Es la edición EE. UU. de `../plan-pizzeria/calculadora_piz_es.py`: MISMAS hojas (`Dough` ·
`Food Cost` · `Instructions`), filas, helpers, `inject_cache` y autotest (pycel + caché = `calcular()`, 3 juegos). Cambios D145:

  (1) masa en GRAMOS (el oficio US trabaja en gramos y % panadero) con eco gris en oz (bola) y lb (masa y harina);
  (2) precios en $/lb con la constante gris «453.6 g per lb»; porciones de `Food Cost` en oz con «16 oz per lb»
      (el gate prohíbe números tecleados en fórmulas salvo 0/1);
  (3) `Food Cost!B5` «Sales tax included in your menu prices (%)» = 0 % (menú US sin impuesto; misma fórmula);
  (4) bloque A «traditional Neapolitan ratios (EU TSG spec, also used by the AVPN)», semáforo «Within / Outside the
      traditional ratios» y aviso fijo: en EE. UU. ninguna norma regula «Neapolitan-style» y cumplir las proporciones NO
      da el «VPN-certified» (certificación privada de la AVPN);
  (5) 8 pizzas (Margherita $18 · Marinara $15 · Diavola $20 · Capricciosa $21 · Quattro Formaggi $21 · Ham & Mushroom $20 ·
      Bufalina $23 · Sausage & Broccoli Rabe $22); ingredientes y caja ($0.75) = «Kit estimate (Oct 2026): replace with your
      distributor's price»;
  (6) US Letter, formatos sin símbolo de moneda (los rótulos llevan «($)», D8), docProps D20.

Independiente de `mapas.py` salvo por los detectores (`restos_espanol`, `no_latinos`) y los textos fijos (firma, aviso, pie).

Corre en el VPS (openpyxl + pycel):

    /root/venv-guias/bin/python calculadora_piz.py --salida <dir>             # genera + inject_cache + autotest
    /root/venv-guias/bin/python calculadora_piz.py --salida <dir> --autotest

Sesión Claude Code, 25-oct-2026.
"""
import argparse
import math
import os
import re
import sys
from collections import OrderedDict

import openpyxl
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill, Protection
from openpyxl.worksheet.datavalidation import DataValidation

AQUI = os.path.dirname(os.path.abspath(__file__))
PD = os.path.dirname(AQUI)
sys.path.insert(0, os.path.join(PD, 'planes-v2_0'))
sys.path.insert(0, os.path.join(PD, 'business-plans-en'))
sys.path.insert(0, PD)
import motor                                                                       # noqa: E402  (solo gate_formatos)
import mapas as M                                                                  # noqa: E402  (detectores y textos fijos)

SLUG = 'pizza-shop-business-plan'
PRODUCTO = 'Pizza Shop Business Plan Kit'
URL = 'aichef.pro/en/digital-products/' + SLUG
FICHERO = 'pizza-dough-food-cost-calculator.xlsx'
TITULO = 'Pizza Dough & Food Cost Calculator (Neapolitan-Style Dough + Food Cost for 8 Pizzas)'
KEYWORDS = ('pizza shop business plan, pizzeria business plan, how to open a pizza shop, pizza shop startup costs, '
            'wood fired pizza oven cost, AI Chef Pro')
VERSION, MES = '1.0', 'October 2026'
VERSION_LINEA = 'Version %s · %s · %s · info@aichef.pro' % (VERSION, MES, URL)
FOOTER = M.FOOTER_EN                                                               # 'AI Chef Pro · aichef.pro · Page &P of &N'
HM, HE, HI = 'Dough', 'Food Cost', 'Instructions'
VERDE = PatternFill('solid', fgColor='E8F5E9')
GRIS = PatternFill('solid', fgColor='F2F2F2')
ROJO = PatternFill('solid', bgColor='F8CBAD', fgColor='F8CBAD')
AMBAR = PatternFill('solid', bgColor='FFE699', fgColor='FFE699')
OK = PatternFill('solid', bgColor='C6E0B4', fgColor='C6E0B4')
FUENTE = 'Calibri'
USD, ENT, PCT, DEC, GR, USD3 = '#,##0.00', '#,##0', '0.0%', '#,##0.0', '#,##0.0', '#,##0.000'
OZ = '#,##0.00'
METODOS = ['Traditional ratios (block A)', "Baker's % (block B)"]
SEM_OK, SEM_NO = 'Within the traditional ratios', 'Outside the traditional ratios'
FC_OK, FC_VIG, FC_ALTO = 'On target', 'Watch the cost', 'Above target'
KIT_PRECIO = "Kit estimate (Oct 2026): replace with your distributor's price"
KIT_LIBRE = "Kit estimate (Oct 2026): free slot for your own pizzas; replace with your distributor's price"

NOTA_VPN = ('FIXED NOTICE. The traffic light on this sheet only tells you whether your flour, salt, yeast and dough ball weight '
            'fall inside the traditional Neapolitan ratios. It does NOT make your pizza "VPN-certified". In the United States '
            'no federal rule governs the term "Neapolitan-style", so you may write it on your menu; but "VPN-certified" is a '
            'private certification awarded by the AVPN after an inspection of your wood-fired dome oven, ingredients and '
            'method, and only a certified shop may use that seal. Meeting these ratios is not that certificate. Until you '
            'hold it, write "Neapolitan-style" and the name of the pizza (Margherita, Marinara...).')

# ---------------------------------------------------------------------------------------------------------------------
# Datos. Todos los precios son «Kit estimate (Oct 2026)» (el research US no trae precios de ingredientes). Sin sales tax.
# ---------------------------------------------------------------------------------------------------------------------
#: bloque A / B: (clave, rótulo, formato, valor por defecto, gris (mín, máx) o None, nota)
ENT_MASA = [
    ('bolas', 'Number of dough balls you want to make', ENT, 60, None,
     'Whole number greater than 0. One dough ball, one pizza.'),
    ('peso', 'Dough ball weight (g)', GR, 250, (180, 250),
     'The traditional Neapolitan ratios call for hand-shaped dough balls of 180 to 250 g. Gray cells show the traditional range.'),
    ('merma', 'Mixing and balling loss (%)', PCT, 0.03, None,
     'Kit estimate: dough left in the mixer, on the bench and on your hands. It is added to the dough you need.'),
    ('h_l', 'Flour per liter of water (g)', GR, 1800, (1600, 1800),
     'Traditional ratios: up to 1,800 g of flour per liter of water (the "kneading point"). The 1,600 g minimum gives a 62 % '
     'absorption (kit calculation over the traditional 55-62 % range).'),
    ('s_l', 'Salt per liter of water (g)', GR, 52.5, (50, 55), 'Traditional ratios: 50 to 55 g of salt per liter of water.'),
    ('l_l', 'Fresh brewer\'s yeast per liter of water (g)', GR, 3, (3, 3), 'Traditional ratios: 3 g per liter of water.'),
    ('agua_l', 'Water per liter (g)', GR, 1000, None, 'Constant: one liter of water weighs 1,000 g. Not editable.'),
    ('g_lb', 'Grams per lb (constant: 453.6 g per lb)', GR, 453.6, None,
     'Constant: 1 lb = 453.6 g. Not editable. It converts the $/lb prices to grams.'),
    ('oz_lb', 'Ounces per lb (constant: 16 oz per lb)', ENT, 16, None,
     'Constant: 1 lb = 16 oz. Not editable. The "Food Cost" sheet uses it for the pizza portions.'),
    ('p_h', '00 flour price ($/lb)', USD, 1.10, None,
     'Kit estimate (Oct 2026): 00 pizza flour in 50 lb bags. Replace with your distributor\'s price.'),
    ('p_s', 'Salt price ($/lb)', USD, 0.60, None, "Kit estimate (Oct 2026): replace with your distributor's price."),
    ('p_l', 'Fresh yeast price ($/lb)', USD, 4.50, None, "Kit estimate (Oct 2026): replace with your distributor's price."),
]
ENT_B = [
    ('hid', 'Hydration (water over flour, %)', PCT, 0.60, None,
     'The traditional ratios leave 55 to 62 % depending on the flour absorption. 60 % is a kit example.'),
    ('sal_b', 'Salt (over flour, %)', '0.00%', 0.028, None, 'Traditional ratios: 2.8 to 3.1 % over the flour (50-55 g per liter of water).'),
    ('lev_b', 'Fresh yeast (over flour, %)', '0.00%', 0.002, None,
     'Traditional ratios: about 0.17 % over the flour (3 g per liter of water). 0.2 % is a kit example.'),
]

# ingredientes: nombre, $/lb, merma %, origen del precio
ING = [
    ('San Marzano-style crushed tomatoes', 2.50, 0.0, KIT_PRECIO),
    ('Fior di latte (fresh mozzarella)', 7.00, 0.03, KIT_PRECIO),
    ('Mozzarella di bufala', 14.00, 0.03, KIT_PRECIO),
    ('Extra virgin olive oil', 6.00, 0.0, KIT_PRECIO),
    ('Fresh basil', 16.00, 0.30, KIT_PRECIO),
    ('Dried oregano', 12.00, 0.0, KIT_PRECIO),
    ('Garlic', 3.50, 0.20, KIT_PRECIO),
    ('Spicy salami', 12.00, 0.03, KIT_PRECIO),
    ('Cooked ham', 7.50, 0.03, KIT_PRECIO),
    ('Sliced mushrooms', 3.50, 0.10, KIT_PRECIO),
    ('Artichoke hearts in oil', 7.00, 0.10, KIT_PRECIO),
    ('Black olives', 5.50, 0.05, KIT_PRECIO),
    ('Gorgonzola', 9.00, 0.03, KIT_PRECIO),
    ('Grana Padano', 11.00, 0.05, KIT_PRECIO),
    ('Smoked provola', 9.50, 0.03, KIT_PRECIO),
    ('Fresh Italian sausage', 5.50, 0.05, KIT_PRECIO),
    ('Broccoli rabe', 4.00, 0.25, KIT_PRECIO),
    ('Cherry tomatoes', 4.50, 0.05, KIT_PRECIO),
    ('Arugula', 7.00, 0.10, KIT_LIBRE),
    ('Prosciutto', 16.00, 0.05, KIT_LIBRE),
    ("'Nduja", 14.00, 0.02, KIT_LIBRE),
    ('Ricotta', 4.50, 0.02, KIT_LIBRE),
]
#: pizza → (menu price, {ingrediente: onzas})
PIZZAS = OrderedDict([
    ('Margherita', (18.00, {'San Marzano-style crushed tomatoes': 3.0, 'Fior di latte (fresh mozzarella)': 4.0,
                            'Extra virgin olive oil': 0.25, 'Fresh basil': 0.1})),
    ('Marinara', (15.00, {'San Marzano-style crushed tomatoes': 4.0, 'Garlic': 0.1, 'Dried oregano': 0.02,
                          'Extra virgin olive oil': 0.25})),
    ('Diavola', (20.00, {'San Marzano-style crushed tomatoes': 3.0, 'Fior di latte (fresh mozzarella)': 3.5, 'Spicy salami': 2.0,
                         'Extra virgin olive oil': 0.25, 'Fresh basil': 0.1})),
    ('Capricciosa', (21.00, {'San Marzano-style crushed tomatoes': 3.0, 'Fior di latte (fresh mozzarella)': 3.0, 'Cooked ham': 1.5,
                             'Sliced mushrooms': 1.5, 'Artichoke hearts in oil': 1.0, 'Black olives': 0.5,
                             'Extra virgin olive oil': 0.25})),
    ('Quattro Formaggi', (21.00, {'Fior di latte (fresh mozzarella)': 2.25, 'Gorgonzola': 1.5, 'Grana Padano': 0.75,
                                  'Smoked provola': 1.5, 'Extra virgin olive oil': 0.15})),
    ('Ham & Mushroom', (20.00, {'San Marzano-style crushed tomatoes': 3.0, 'Fior di latte (fresh mozzarella)': 3.5, 'Cooked ham': 1.75,
                                'Sliced mushrooms': 1.75, 'Extra virgin olive oil': 0.25})),
    ('Bufalina', (23.00, {'San Marzano-style crushed tomatoes': 2.0, 'Mozzarella di bufala': 4.5, 'Cherry tomatoes': 1.0,
                          'Fresh basil': 0.1, 'Extra virgin olive oil': 0.25})),
    ('Sausage & Broccoli Rabe', (22.00, {'Smoked provola': 2.5, 'Fresh Italian sausage': 3.0, 'Broccoli rabe': 2.5, 'Garlic': 0.1,
                                        'Extra virgin olive oil': 0.25})),
])
NP, NI = len(PIZZAS), len(ING)
PCOL = [openpyxl.utils.get_column_letter(5 + j) for j in range(NP)]               # E..L
ESC_IN = [('tax', 'Sales tax included in your menu prices (%)', PCT, 0.0,
           'Leave 0 % if your menu prices are quoted before sales tax (the usual way in the US: the guest pays tax on top). '
           'If your prices already include tax, enter the rate so the food cost is calculated on the price net of tax.'),
          ('obj', 'Target food cost (% of the pre-tax menu price)', PCT, 0.26,
           'Industry reference for pizza shops: 25-30 %. The kit\'s business plan uses 26 %. Use the one in your financial plan.'),
          ('tol', 'Watch margin above the target (points)', PCT, 0.05,
           'Up to this margin above the target, the traffic light says "Watch the cost"; beyond it, "Above target".'),
          ('caja', 'Cost of the take-out box ($)', USD, 0.75,
           "Kit estimate (Oct 2026): replace with your distributor's price."),
          ]

# ---------------------------------------------------------------------------------------------------------------------
# Disposición de la hoja Dough
# ---------------------------------------------------------------------------------------------------------------------
M_ = OrderedDict()                                                                 # clave → fila en Dough
_f_ = 5
M_HA = _f_
for k, *_r in ENT_MASA:
    _f_ += 1
    M_[k] = _f_
M_RA = _f_ + 2
ECO = ('peso_oz', 'masa_lb', 'harina_lb')                                          # filas grises con el eco en oz / lb
SAL_A = [
    ('peso_oz', 'Dough ball weight (oz, echo of the ball weight above)', OZ, '{peso}*{oz_lb}/{g_lb}',
     'Gray echo of the green cell in ounces: grams × 16 ÷ 453.6.'),
    ('masa', 'Total dough to prepare (g)', GR, '{bolas}*{peso}*(1+{merma})', 'Dough balls × weight × (1 + loss).'),
    ('masa_lb', 'Total dough to prepare (lb, echo)', OZ, '{masa}/{g_lb}', 'Gray echo in pounds.'),
    ('masa_l', 'Dough that comes out of each liter of water (g)', GR, '{agua_l}+{h_l}+{s_l}+{l_l}',
     'Water + flour + salt + yeast per liter.'),
    ('litros', 'Liters of water you need', '#,##0.00', '{masa}/{masa_l}', 'Total dough ÷ dough per liter.'),
    ('agua', 'Water (g)', GR, '{litros}*{agua_l}', ''),
    ('harina', 'Flour (g)', GR, '{litros}*{h_l}', ''),
    ('harina_lb', 'Flour (lb, echo)', OZ, '{harina}/{g_lb}', 'Gray echo in pounds: compare it with your 50 lb bags.'),
    ('sal', 'Salt (g)', GR, '{litros}*{s_l}', ''),
    ('lev', 'Fresh yeast (g)', GR, '{litros}*{l_l}', ''),
    ('ph_h', "Baker's %: hydration (water ÷ flour)", PCT, '{agua}/{harina}', 'To carry it over to block B.'),
    ('ph_s', "Baker's %: salt (salt ÷ flour)", '0.00%', '{sal}/{harina}', ''),
    ('ph_l', "Baker's %: yeast (yeast ÷ flour)", '0.00%', '{lev}/{harina}', ''),
    ('h_bola', 'Flour per dough ball (g)', GR, '{harina}/{bolas}', ''),
    ('bolas_lb', 'Dough balls that come out of 1 lb of flour', '#,##0.0', '{bolas}*{g_lb}/{harina}',
     'With the traditional ratios, about 2.8 balls of 250 g per lb of flour.'),
    ('coste_a', 'Dough cost ($)', USD, '({harina}*{p_h}+{sal}*{p_s}+{lev}*{p_l})/{g_lb}',
     'Sum of grams × $/lb ÷ 453.6, loss included.'),
    ('cbola_a', 'Dough cost per ball ($)', USD3, '{coste_a}/{bolas}', ''),
    ('sem', 'Traditional ratios traffic light', 'General',
     'IF(AND({h_l}>={h_min},{h_l}<={h_max},{s_l}>={s_min},{s_l}<={s_max},{peso}>={p_min},{peso}<={p_max}),"%s","%s")' % (SEM_OK, SEM_NO),
     'It looks at flour and salt per liter and at the ball weight. Mind the fixed notice above: it is not a seal.'),
]
S_A = OrderedDict((k, M_RA + 1 + i) for i, (k, *_r) in enumerate(SAL_A))
M_HB = M_RA + 1 + len(SAL_A) + 1
for i, (k, *_r) in enumerate(ENT_B):
    M_[k] = M_HB + 1 + i
M_RB = M_HB + 1 + len(ENT_B) + 1
SAL_B = [
    ('harina_b', 'Flour (g)', GR, '{masa}/(1+{hid}+{sal_b}+{lev_b})', 'Dough ÷ (1 + hydration + salt + yeast), as decimals.'),
    ('harina_b_lb', 'Flour (lb, echo)', OZ, '{harina_b}/{g_lb}', 'Gray echo in pounds.'),
    ('agua_b', 'Water (g)', GR, '{harina_b}*{hid}', ''),
    ('sal_bg', 'Salt (g)', GR, '{harina_b}*{sal_b}', ''),
    ('lev_bg', 'Fresh yeast (g)', GR, '{harina_b}*{lev_b}', ''),
    ('coste_b', 'Dough cost ($)', USD, '({harina_b}*{p_h}+{sal_bg}*{p_s}+{lev_bg}*{p_l})/{g_lb}', 'Same $/lb as block A.'),
    ('cbola_b', 'Dough cost per ball ($)', USD3, '{coste_b}/{bolas}', ''),
    ('bolas_lb_b', 'Dough balls that come out of 1 lb of flour', '#,##0.0', '{bolas}*{g_lb}/{harina_b}', ''),
]
ECO_B = ('harina_b_lb',)
S_B = OrderedDict((k, M_RB + 1 + i) for i, (k, *_r) in enumerate(SAL_B))
M_SEL = M_RB + 1 + len(SAL_B) + 1
M_['metodo'] = M_SEL + 1
S_M = {'cbola': M_SEL + 2}
M_FIN = M_SEL + 3

# ---------------------------------------------------------------------------------------------------------------------
# Disposición de la hoja Food Cost
# ---------------------------------------------------------------------------------------------------------------------
E = OrderedDict((k, 5 + i) for i, (k, *_r) in enumerate(ESC_IN))
E['oz_lb'] = 5 + len(ESC_IN)                                                       # constante gris (enlace a Dough)
E_CM = E['oz_lb'] + 1                                                              # coste de la masa por bola (enlace)
E_CAB = E_CM + 2
E_PVP = E_CAB + 1
E_I0 = E_PVP + 1
E_I1 = E_I0 + NI - 1
E_R = E_I1 + 2
RES = [
    ('peso_t', 'Raw weight of the pizza (oz)', OZ),
    ('c_masa', 'Dough cost ($)', USD3),
    ('c_ing', 'Ingredient cost ($)', USD3),
    ('c_tot', 'COST OF THE PIZZA ($)', USD3),
    ('pvp_sin', 'Menu price before sales tax ($)', USD),
    ('fc', 'FOOD COST (%)', PCT),
    ('sem', 'Traffic light against the target', 'General'),
    ('margen', 'Margin per pizza ($, price before tax − cost)', USD),
    ('pvp_obj', 'Menu price needed to reach the target ($)', USD),
    ('c_llevar', 'Take-out cost: pizza + box ($)', USD3),
    ('fc_llevar', 'Take-out food cost with the box (%)', PCT),
]
R = OrderedDict((k, E_R + 1 + i) for i, (k, *_r) in enumerate(RES))
E_FIN = E_R + len(RES) + 1


def valores_defecto():
    """Todas las entradas por defecto en un solo dict plano."""
    v = OrderedDict()
    for k, _r, _f2, val, _g, _n in ENT_MASA + ENT_B:
        v[k] = val
    v['metodo'] = METODOS[0]
    for k, _r, _f2, val, _n in ESC_IN:
        v[k] = val
    for i, (nom, p, m, _o) in enumerate(ING):
        v['p%d' % i], v['m%d' % i] = p, m
    for j, (nom, (pvp, gr)) in enumerate(PIZZAS.items()):
        v['pvp%d' % j] = pvp
        for i, (ni, *_r) in enumerate(ING):
            v['g%d_%d' % (i, j)] = float(gr.get(ni, 0))
    return v


# ---------------------------------------------------------------------------------------------------------------------
# Cálculo en Python (espejo de las fórmulas) — fuente del autotest
# ---------------------------------------------------------------------------------------------------------------------
GRIS_MM = dict(h_min=1600, h_max=1800, s_min=50, s_max=55, p_min=180, p_max=250)


def calcular(v):
    r = OrderedDict()
    r['peso_oz'] = v['peso'] * v['oz_lb'] / v['g_lb']
    r['masa'] = v['bolas'] * v['peso'] * (1 + v['merma'])
    r['masa_lb'] = r['masa'] / v['g_lb']
    r['masa_l'] = v['agua_l'] + v['h_l'] + v['s_l'] + v['l_l']
    r['litros'] = r['masa'] / r['masa_l']
    r['agua'] = r['litros'] * v['agua_l']
    r['harina'] = r['litros'] * v['h_l']
    r['harina_lb'] = r['harina'] / v['g_lb']
    r['sal'] = r['litros'] * v['s_l']
    r['lev'] = r['litros'] * v['l_l']
    r['ph_h'] = r['agua'] / r['harina']
    r['ph_s'] = r['sal'] / r['harina']
    r['ph_l'] = r['lev'] / r['harina']
    r['h_bola'] = r['harina'] / v['bolas']
    r['bolas_lb'] = v['bolas'] * v['g_lb'] / r['harina']
    r['coste_a'] = (r['harina'] * v['p_h'] + r['sal'] * v['p_s'] + r['lev'] * v['p_l']) / v['g_lb']
    r['cbola_a'] = r['coste_a'] / v['bolas']
    g = GRIS_MM
    r['sem'] = SEM_OK if (g['h_min'] <= v['h_l'] <= g['h_max'] and g['s_min'] <= v['s_l'] <= g['s_max']
                          and g['p_min'] <= v['peso'] <= g['p_max']) else SEM_NO
    r['harina_b'] = r['masa'] / (1 + v['hid'] + v['sal_b'] + v['lev_b'])
    r['harina_b_lb'] = r['harina_b'] / v['g_lb']
    r['agua_b'] = r['harina_b'] * v['hid']
    r['sal_bg'] = r['harina_b'] * v['sal_b']
    r['lev_bg'] = r['harina_b'] * v['lev_b']
    r['coste_b'] = (r['harina_b'] * v['p_h'] + r['sal_bg'] * v['p_s'] + r['lev_bg'] * v['p_l']) / v['g_lb']
    r['cbola_b'] = r['coste_b'] / v['bolas']
    r['bolas_lb_b'] = v['bolas'] * v['g_lb'] / r['harina_b']
    r['cbola'] = r['cbola_a'] if v['metodo'] == METODOS[0] else r['cbola_b']
    neto = [v['p%d' % i] / (1 - v['m%d' % i]) for i in range(NI)]
    r['neto'] = neto
    for j in range(NP):
        gs = [v['g%d_%d' % (i, j)] for i in range(NI)]
        p = OrderedDict()
        p['peso_t'] = r['peso_oz'] + sum(gs)
        p['c_masa'] = r['cbola']
        p['c_ing'] = sum(g_ * n for g_, n in zip(gs, neto)) / v['oz_lb']
        p['c_tot'] = p['c_masa'] + p['c_ing']
        p['pvp_sin'] = v['pvp%d' % j] / (1 + v['tax'])
        p['fc'] = p['c_tot'] / p['pvp_sin']
        p['sem'] = FC_OK if p['fc'] <= v['obj'] else (FC_VIG if p['fc'] <= v['obj'] + v['tol'] else FC_ALTO)
        p['margen'] = p['pvp_sin'] - p['c_tot']
        p['pvp_obj'] = p['c_tot'] / v['obj'] * (1 + v['tax'])
        p['c_llevar'] = p['c_tot'] + v['caja']
        p['fc_llevar'] = p['c_llevar'] / p['pvp_sin']
        r['p%d' % j] = p
    return r


# ---------------------------------------------------------------------------------------------------------------------
# Helpers (de la versión ES)
# ---------------------------------------------------------------------------------------------------------------------
def _f(ws, coord, v, b=False, sz=11, i=False, fmt=None, wrap=False, fill=None, locked=True):
    c = ws[coord]
    c.value = v
    c.font = Font(name=FUENTE, size=sz, bold=b, italic=i)
    if fmt:
        c.number_format = fmt
    c.alignment = Alignment(wrap_text=wrap, vertical='top')
    if fill:
        c.fill = fill
    c.protection = Protection(locked=locked)
    return c


def _nota(ws, fila, texto, desde='E', hasta='F', chars=80.0):
    if not texto:
        return
    _f(ws, '%s%d' % (desde, fila), texto, wrap=True)
    ws.merge_cells('%s%d:%s%d' % (desde, fila, hasta, fila))
    ws.row_dimensions[fila].height = max(16, 15 * math.ceil(len(texto) / chars))


def _imprimir(ws, horizontal=False):
    ws.page_setup.paperSize = M.PAPEL_EN                                         # 1 = US Letter
    ws.page_setup.orientation = 'landscape' if horizontal else 'portrait'
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.oddFooter.center.text = FOOTER
    ws.print_options.horizontalCentered = True
    ws.protection.sheet = True                                                    # sin contraseña
    ws.sheet_view.showGridLines = False


def _dv(ws, tipo, celdas, f1=None, f2=None, op=None, titulo='Invalid value', msg=''):
    dv = DataValidation(type=tipo, formula1=f1, formula2=f2, operator=op, allow_blank=False, showErrorMessage=True,
                        errorTitle=titulo, error=msg)
    ws.add_data_validation(dv)
    for c in celdas:
        dv.add(c)


def a_excel(expr, ref):
    return re.sub(r'\{(\w+)\}', lambda m: ref[m.group(1)], expr)


def refs_masa():
    ref = dict((k, '$B$%d' % f) for k, f in M_.items())
    ref.update((k, '$B$%d' % f) for k, f in S_A.items())
    ref.update((k, '$B$%d' % f) for k, f in S_B.items())
    ref['p_min'], ref['p_max'] = '$C$%d' % M_['peso'], '$D$%d' % M_['peso']
    ref['h_min'], ref['h_max'] = '$C$%d' % M_['h_l'], '$D$%d' % M_['h_l']
    ref['s_min'], ref['s_max'] = '$C$%d' % M_['s_l'], '$D$%d' % M_['s_l']
    return ref


# ---------------------------------------------------------------------------------------------------------------------
# Hojas
# ---------------------------------------------------------------------------------------------------------------------
def hoja_masa(wb):
    ws = wb.active
    ws.title = HM
    for col, w in (('A', 58), ('B', 17), ('C', 12), ('D', 12), ('E', 48), ('F', 30)):
        ws.column_dimensions[col].width = w
    _f(ws, 'A1', 'PIZZA DOUGH CALCULATOR: TRADITIONAL NEAPOLITAN RATIOS AND BAKER\'S %', b=True, sz=14)
    _f(ws, 'A2', 'Type only in the GREEN cells: everything else is calculated. Dough in grams (how bakers work), with a gray echo in '
                 'ounces and pounds. Prices in $/lb. The ingredient prices are examples (Kit estimate, Oct 2026): replace them with '
                 'your distributor\'s.', wrap=True)
    ws.merge_cells('A2:F2')
    ws.row_dimensions[2].height = 32
    _f(ws, 'A3', NOTA_VPN, wrap=True, b=True, fill=GRIS)
    ws.merge_cells('A3:F3')
    ws.row_dimensions[3].height = 15 * math.ceil(len(NOTA_VPN) / 150.0) + 6
    for col, t in zip('ABCDE', ('A. TRADITIONAL NEAPOLITAN RATIOS (EU TSG spec, also used by the AVPN), PER LITER OF WATER', 'Value',
                                'Trad. min.', 'Trad. max.', 'Notes')):
        _f(ws, '%s%d' % (col, M_HA), t, b=True, fill=GRIS, wrap=True)
    ws.row_dimensions[M_HA].height = 32
    for k, rot, fmt, val, gris, nota in ENT_MASA:
        f = M_[k]
        _f(ws, 'A%d' % f, rot)
        if k in ('agua_l', 'g_lb', 'oz_lb'):
            _f(ws, 'B%d' % f, val, fmt=fmt, fill=GRIS)
        else:
            _f(ws, 'B%d' % f, val, fmt=fmt, fill=VERDE, locked=False)
        if gris:
            _f(ws, 'C%d' % f, gris[0], fmt=fmt, fill=GRIS)
            _f(ws, 'D%d' % f, gris[1], fmt=fmt, fill=GRIS)
        _nota(ws, f, nota)
    ref = refs_masa()
    for col, t in zip('ABE', ('RESULT OF BLOCK A', 'Value', 'Notes')):
        _f(ws, '%s%d' % (col, M_RA), t, b=True, fill=GRIS)
    for k, rot, fmt, expr, nota in SAL_A:
        f = S_A[k]
        neg = k in ('masa', 'cbola_a', 'sem')
        eco = k in ECO
        _f(ws, 'A%d' % f, rot, b=neg, i=eco, fill=GRIS if eco else None)
        _f(ws, 'B%d' % f, '=' + a_excel(expr, ref), fmt=fmt, b=neg, wrap=(k == 'sem'), i=eco, fill=GRIS if eco else None)
        _nota(ws, f, nota)
    ws.row_dimensions[S_A['sem']].height = 48
    for col, t in zip('ABE', ("B. BY BAKER'S % (over the flour)", 'Value', 'Notes')):
        _f(ws, '%s%d' % (col, M_HB), t, b=True, fill=GRIS)
    for k, rot, fmt, val, gris, nota in ENT_B:
        f = M_[k]
        _f(ws, 'A%d' % f, rot)
        _f(ws, 'B%d' % f, val, fmt=fmt, fill=VERDE, locked=False)
        _nota(ws, f, nota)
    for col, t in zip('ABE', ('RESULT OF BLOCK B (same total dough as block A)', 'Value', 'Notes')):
        _f(ws, '%s%d' % (col, M_RB), t, b=True, fill=GRIS)
    for k, rot, fmt, expr, nota in SAL_B:
        f = S_B[k]
        neg = k == 'cbola_b'
        eco = k in ECO_B
        _f(ws, 'A%d' % f, rot, b=neg, i=eco, fill=GRIS if eco else None)
        _f(ws, 'B%d' % f, '=' + a_excel(expr, ref), fmt=fmt, b=neg, i=eco, fill=GRIS if eco else None)
        _nota(ws, f, nota)
    _f(ws, 'A%d' % M_SEL, 'DOUGH COST USED BY THE FOOD COST SHEET', b=True, fill=GRIS)
    _f(ws, 'B%d' % M_SEL, 'Value', b=True, fill=GRIS)
    _f(ws, 'E%d' % M_SEL, 'Notes', b=True, fill=GRIS)
    f = M_['metodo']
    _f(ws, 'A%d' % f, 'Block that feeds the Food Cost sheet')
    _f(ws, 'B%d' % f, METODOS[0], fill=VERDE, locked=False, wrap=True)
    ws.row_dimensions[f].height = 32
    _nota(ws, f, 'Choose which dough cost per ball the "%s" sheet uses: the one from the traditional ratios (A) or the one from your '
                 'own percentages (B).' % HE)
    f = S_M['cbola']
    _f(ws, 'A%d' % f, 'Dough cost per ball for the Food Cost sheet ($)', b=True)
    _f(ws, 'B%d' % f, '=IF(%s="%s",%s,%s)' % ('$B$%d' % M_['metodo'], METODOS[0], '$B$%d' % S_A['cbola_a'], '$B$%d' % S_B['cbola_b']),
       fmt=USD3, b=True)
    # semáforo en color
    sem = '$B$%d' % S_A['sem']
    ws.conditional_formatting.add('B%d' % S_A['sem'], FormulaRule(formula=['%s="%s"' % (sem, SEM_OK)], fill=OK))
    ws.conditional_formatting.add('B%d' % S_A['sem'], FormulaRule(formula=['%s="%s"' % (sem, SEM_NO)], fill=ROJO))
    # validaciones
    cel = lambda k: 'B%d' % M_[k]
    _dv(ws, 'whole', [cel('bolas')], f1='1', f2='100000', op='between', titulo='Invalid number',
        msg='Enter a whole number of 1 or more.')
    _dv(ws, 'decimal', [cel(k) for k in ('peso', 'h_l', 's_l', 'l_l', 'p_h', 'p_s', 'p_l')], f1='0', op='greaterThan',
        titulo='Invalid value', msg='Enter a number greater than 0.')
    _dv(ws, 'decimal', [cel('merma'), cel('sal_b'), cel('lev_b')], f1='0', f2='1', op='between', titulo='Invalid percentage',
        msg='Enter a percentage between 0 and 1 (0.03 = 3 %).')
    _dv(ws, 'decimal', [cel('hid')], f1='0', f2='2', op='between', titulo='Invalid hydration',
        msg='Enter a percentage between 0 and 200 % (0.6 = 60 %).')
    _dv(ws, 'list', [cel('metodo')], f1='"%s"' % ','.join(METODOS), titulo='Invalid option', msg='Choose an option from the list.')
    _imprimir(ws)


def hoja_escandallo(wb):
    ws = wb.create_sheet(HE)
    for col, w in (('A', 42), ('B', 13), ('C', 10), ('D', 13)):
        ws.column_dimensions[col].width = w
    for col in PCOL:
        ws.column_dimensions[col].width = 14
    ws.column_dimensions['M'].width = 62
    _f(ws, 'A1', 'FOOD COST PER PIZZA: 8 EXAMPLE PIZZAS, PORTIONS IN OUNCES', b=True, sz=14)
    _f(ws, 'A2', 'Type only in the GREEN cells (ingredients, prices, waste, ounces, names and menu prices). Prices are $/lb before sales '
                 'tax; portions are in ounces. The prices are examples (Kit estimate, Oct 2026): use your distributor\'s. The sample '
                 'menu says "Neapolitan-style" and uses the recipe names, not a "VPN-certified" claim (see the notice on the "%s" '
                 'sheet).' % HM, wrap=True)
    ws.merge_cells('A2:M2')
    ws.row_dimensions[2].height = 32
    for col, t in zip('ABE', ('GENERAL DATA', 'Value', 'Notes')):
        _f(ws, '%s4' % col, t, b=True, fill=GRIS)
    for k, rot, fmt, val, nota in ESC_IN:
        f = E[k]
        _f(ws, 'A%d' % f, rot)
        _f(ws, 'B%d' % f, val, fmt=fmt, fill=VERDE, locked=False)
        _nota(ws, f, nota, desde='E', hasta='M', chars=140.0)
    f = E['oz_lb']
    _f(ws, 'A%d' % f, 'Ounces per lb (constant: 16 oz per lb)', i=True, fill=GRIS)
    _f(ws, 'B%d' % f, "='%s'!$B$%d" % (HM, M_['oz_lb']), fmt=ENT, i=True, fill=GRIS)
    _nota(ws, f, 'Constant, linked to the "%s" sheet. Not editable. It converts the $/lb prices to the ounce portions.' % HM,
          desde='E', hasta='M', chars=140.0)
    _f(ws, 'A%d' % E_CM, 'Dough cost per ball (from the "%s" sheet) ($)' % HM, b=True)
    _f(ws, 'B%d' % E_CM, "='%s'!$B$%d" % (HM, S_M['cbola']), fmt=USD3, b=True)
    _nota(ws, E_CM, 'Calculated on the "%s" sheet (block A or B, as you choose there).' % HM, desde='E', hasta='M', chars=140.0)
    # cabecera de la tabla
    for col, t in (('A', 'INGREDIENT (portions in oz)'), ('B', '$/lb before tax'), ('C', 'Waste'), ('D', 'Net $/lb')):
        _f(ws, '%s%d' % (col, E_CAB), t, b=True, fill=GRIS, wrap=True)
    _f(ws, 'M%d' % E_CAB, 'ORIGIN OF THE EXAMPLE PRICE', b=True, fill=GRIS, wrap=True)
    _f(ws, 'A%d' % E_PVP, 'Menu price ($)', b=True)
    for j, (nom, (pvp, _gr)) in enumerate(PIZZAS.items()):
        col = PCOL[j]
        _f(ws, '%s%d' % (col, E_CAB), nom, b=True, fill=VERDE, wrap=True, locked=False)
        _f(ws, '%s%d' % (col, E_PVP), pvp, fmt=USD, fill=VERDE, locked=False, b=True)
    ws.row_dimensions[E_CAB].height = 32
    for i, (nom, p, m, origen) in enumerate(ING):
        f = E_I0 + i
        _f(ws, 'A%d' % f, nom, fill=VERDE, locked=False)
        _f(ws, 'B%d' % f, p, fmt=USD, fill=VERDE, locked=False)
        _f(ws, 'C%d' % f, m, fmt='0%', fill=VERDE, locked=False)
        _f(ws, 'D%d' % f, '=B%d/(1-C%d)' % (f, f), fmt=USD)
        for j, (pn, (_pvp, gr)) in enumerate(PIZZAS.items()):
            _f(ws, '%s%d' % (PCOL[j], f), float(gr.get(nom, 0)), fmt='#,##0.00;-#,##0.00;"-"', fill=VERDE, locked=False)
        _f(ws, 'M%d' % f, origen, i=True)
    _f(ws, 'A%d' % (E_R), 'RESULT PER PIZZA (ounces in the columns above)', b=True, fill=GRIS)
    for j in range(NP):
        _f(ws, '%s%d' % (PCOL[j], E_R), '=%s%d' % (PCOL[j], E_CAB), b=True, fill=GRIS, wrap=True)
    ws.row_dimensions[E_R].height = 32
    ozl = '$B$%d' % E['oz_lb']
    for k, rot, fmt in RES:
        f = R[k]
        neg = k in ('c_tot', 'fc')
        _f(ws, 'A%d' % f, rot, b=neg)
        for j in range(NP):
            c = PCOL[j]
            rng = '%s$%d:%s$%d' % (c, E_I0, c, E_I1)
            ex = {
                'peso_t': "'%s'!$B$%d+SUM(%s)" % (HM, S_A['peso_oz'], rng),
                'c_masa': '$B$%d' % E_CM,
                'c_ing': 'SUMPRODUCT(%s,$D$%d:$D$%d)/%s' % (rng, E_I0, E_I1, ozl),
                'c_tot': '%s%d+%s%d' % (c, R['c_masa'], c, R['c_ing']),
                'pvp_sin': '%s%d/(1+$B$%d)' % (c, E_PVP, E['tax']),
                'fc': '%s%d/%s%d' % (c, R['c_tot'], c, R['pvp_sin']),
                'sem': 'IF(%s%d<=$B$%d,"%s",IF(%s%d<=$B$%d+$B$%d,"%s","%s"))' % (
                    c, R['fc'], E['obj'], FC_OK, c, R['fc'], E['obj'], E['tol'], FC_VIG, FC_ALTO),
                'margen': '%s%d-%s%d' % (c, R['pvp_sin'], c, R['c_tot']),
                'pvp_obj': '%s%d/$B$%d*(1+$B$%d)' % (c, R['c_tot'], E['obj'], E['tax']),
                'c_llevar': '%s%d+$B$%d' % (c, R['c_tot'], E['caja']),
                'fc_llevar': '%s%d/%s%d' % (c, R['c_llevar'], c, R['pvp_sin']),
            }[k]
            _f(ws, '%s%d' % (c, f), '=' + ex, fmt=fmt, b=neg, wrap=(k == 'sem'))
    ws.row_dimensions[R['sem']].height = 32
    rng = '%s%d:%s%d' % (PCOL[0], R['sem'], PCOL[-1], R['sem'])
    top = '%s%d' % (PCOL[0], R['sem'])
    for txt, fill in ((FC_OK, OK), (FC_VIG, AMBAR), (FC_ALTO, ROJO)):
        ws.conditional_formatting.add(rng, FormulaRule(formula=['%s="%s"' % (top, txt)], fill=fill))
    f = E_FIN
    _f(ws, 'A%d' % f, 'The food cost in the table is for the pizza served in the dining room. For take-out the box is added to the cost '
                      'and the menu price is the same; if you charge more for take-out or delivery, change the pizza price to see the '
                      'effect. It does not include the commission of the third-party delivery marketplace, which comes out of the '
                      'price, not out of the cost of the pizza.', i=True, wrap=True)
    ws.merge_cells('A%d:M%d' % (f, f))
    ws.row_dimensions[f].height = 34
    # validaciones
    _dv(ws, 'decimal', ['B%d' % E['tax'], 'B%d' % E['obj'], 'B%d' % E['tol']] + ['C%d' % (E_I0 + i) for i in range(NI)],
        f1='0', f2='0.9', op='between', titulo='Invalid percentage', msg='Enter a percentage between 0 and 0.9 (0.1 = 10 %).')
    _dv(ws, 'decimal', ['B%d' % E['caja']] + ['B%d' % (E_I0 + i) for i in range(NI)]
        + ['%s%d' % (c, E_PVP) for c in PCOL] + ['%s%d' % (c, E_I0 + i) for c in PCOL for i in range(NI)],
        f1='0', op='greaterThanOrEqual', titulo='Invalid amount', msg='Enter a number greater than or equal to 0.')
    _dv(ws, 'decimal', ['B%d' % E['obj']], f1='0.01', f2='0.9', op='between', titulo='Invalid target',
        msg='The target cannot be 0 (the sheet divides by it): enter between 0.01 and 0.9.')
    _imprimir(ws, horizontal=True)


def _instrucciones():
    c = calcular(valores_defecto())
    fc_m = c['p0']['fc'] * 100
    fc_b = c['p%d' % (list(PIZZAS).index('Bufalina'))]['fc'] * 100
    return [
        ('INSTRUCTIONS: Pizza Dough & Food Cost Calculator', True),
        ('Two tools in one workbook: the "%s" sheet tells you how much flour, water, salt and yeast you need for the number of dough '
         'balls you are going to make, and what each ball costs; the "%s" sheet calculates the cost of 8 example pizzas, their food '
         'cost, their margin and the menu price you need. Type only in the GREEN cells; everything else is calculated. Dough is in '
         'grams (how bakers work) with a gray echo in ounces and pounds; ingredient prices are in $/lb and the pizza portions are '
         'in ounces.' % (HM, HE), False),
        ('NOTICE ABOUT THE NAME', True),
        (NOTA_VPN, False),
        ('"%s" SHEET' % HM.upper(), True),
        ('1. Block A (per liter of water): type how many dough balls you want, their weight and the mixing loss. The amounts per '
         'liter are the traditional Neapolitan ratios (the EU Traditional Specialty Guaranteed spec, also used by the AVPN): up to '
         '1,800 g of flour, 50 to 55 g of salt and 3 g of fresh brewer\'s yeast per liter of water, with dough balls of 180 to 250 g. '
         'The traditional range appears in gray next to each cell.', False),
        ('2. The block A traffic light says whether the flour, the salt and the ball weight fit those ranges. Green does not mean '
         '"VPN-certified": that certificate comes from the AVPN after an inspection of the shop, the oven and the method.', False),
        ('3. Block B (baker\'s %): type the hydration, the salt and the yeast as a percentage of the flour weight, which is how the '
         'trade works. It uses the same total dough as block A. If you type the percentages that block A calculates, you get the '
         'same result.', False),
        ('4. At the bottom you choose which dough cost per ball feeds the food cost: the one from the traditional ratios (A) or the '
         'one from your own percentages (B).', False),
        ('"%s" SHEET' % HE.upper(), True),
        ('1. Type the sales tax included in your menu prices (0 % if your menu is quoted before tax, which is the usual US way), the '
         'target food cost, the watch margin and the cost of the take-out box.', False),
        ('2. In the table, each row is an ingredient with its price per lb before tax and its waste (the part lost when trimming or '
         'cleaning). The net price per lb divides the price by (1 − waste). The ounces of each pizza go in its column; you can change '
         'the names, the ounces and the menu prices, and use the free rows for your own ingredients. The conversion from ounces to '
         'pounds uses the gray constant (16 oz per lb).', False),
        ('3. For each pizza you see the cost (dough + ingredients), the price before tax, the food cost with a traffic light, the '
         'margin per pizza, the menu price you would need to reach your target and the take-out food cost with the box.', False),
        ('4. The example Margherita comes out at about %.0f %% food cost. With buffalo mozzarella, Italian cured meats or ham it goes '
         'up: the Bufalina comes out at about %.0f %%. If a pizza lands above the target, raise the price, cut the portion or change '
         'an ingredient.' % (round(fc_m), round(fc_b)), False),
        ('PRICES AND SOURCES', True),
        ('Every default price is an example: "Kit estimate (Oct 2026): replace with your distributor\'s price" (the kit\'s research did '
         'not find reliable distributor prices for these ingredients). Replace them with your distributor\'s before deciding on '
         'prices. The sample menu prices (Marinara $15 to Bufalina $23, Margherita $18) come from the kit\'s business plan case, based '
         'on the Margherita prices of US pizzerias read in October 2026 (about $18).', False),
        ('Sales tax is not part of the menu prices in this sample (0 %). The rates vary by state and city; ask your accountant. Beer, '
         'wine and other drinks are not in this food cost.', False),
        ('Sources: traditional Neapolitan pizza specification (EU Traditional Specialty Guaranteed, followed by the AVPN); US pizza '
         'shop research for the kit (October 2026).', False),
    ]


def hoja_instrucciones(wb):
    ws = wb.create_sheet(HI)
    ws.column_dimensions['A'].width = 110
    f = 1
    for t, b in _instrucciones():
        _f(ws, 'A%d' % f, t, b=b, sz=13 if f == 1 else 11, wrap=True)
        if not b:
            ws.row_dimensions[f].height = max(16, 15 * math.ceil(len(t) / 120.0))
        f += 2 if b and f > 1 else 1
    f += 1
    for t in ('To edit a cell that is not green: Review → Unprotect Sheet (there is no password).',
              'More templates and kits in the catalog at ' + M.URL_TIENDA, M.FIRMA_EN, VERSION_LINEA, M.AVISO_EN):
        _f(ws, 'A%d' % f, t, wrap=True, i=t == M.AVISO_EN)
        f += 1
    _imprimir(ws)


def docprops(wb):
    pr = wb.properties
    pr.title = '%s · %s' % (TITULO, PRODUCTO)
    pr.subject = '%s · v%s' % (PRODUCTO, VERSION)
    pr.keywords = KEYWORDS
    pr.description = URL
    pr.category = 'AI Chef Pro · Digital products'
    pr.creator = pr.lastModifiedBy = 'AI Chef Pro'


def generar(destino):
    wb = openpyxl.Workbook()
    hoja_masa(wb)
    hoja_escandallo(wb)
    hoja_instrucciones(wb)
    docprops(wb)
    os.makedirs(os.path.dirname(destino), exist_ok=True)
    wb.save(destino)
    import inject_cache as IC                                                      # noqa: E402
    nf, inj, fail = IC.inject(destino)
    print('%s: formulas=%d cache_inyectado=%d fallos_pycel=%d' % (os.path.basename(destino), nf, inj, fail))
    return fail


# ---------------------------------------------------------------------------------------------------------------------
# Autotest
# ---------------------------------------------------------------------------------------------------------------------
#: «authentic»/«verace» como promesa, euros, español de la hoja, restos de otros productos y de la reglamentación ES (RX_PIZ_NO)
PROHIBIDOS = [r'authentic', r'\bverace\b', r'€', r'97/2010', r'Napoletana', r'\bIVA\b',
              r'tapas|small plates|cocktails?|spirits|neighborhood bar|barback|last call|happy hour|draft beer system',
              r'\bETG\b|pliego|convenio|\bSMI\b|aut[oó]nomo|Veri\*?factu|SGAE|\bIAE\b|terraza|comensal',
              r'\bfreelance', r'carnet', r'tarifa plana', r'RD 393', r'gluten[- ]free', r'\b15-25 ?%']
CELDA = {}                                                                         # clave de entrada → (hoja, celda)
for _k, _f2 in M_.items():
    CELDA[_k] = (HM, 'B%d' % _f2)
for _k, _f2 in E.items():
    CELDA[_k] = (HE, 'B%d' % _f2)
for _i in range(NI):
    CELDA['p%d' % _i] = (HE, 'B%d' % (E_I0 + _i))
    CELDA['m%d' % _i] = (HE, 'C%d' % (E_I0 + _i))
    for _j in range(NP):
        CELDA['g%d_%d' % (_i, _j)] = (HE, '%s%d' % (PCOL[_j], E_I0 + _i))
for _j in range(NP):
    CELDA['pvp%d' % _j] = (HE, '%s%d' % (PCOL[_j], E_PVP))


def celdas_salida(r):
    """(hoja, celda, valor esperado, etiqueta) de todas las salidas."""
    out = []
    for k in S_A:
        out.append((HM, 'B%d' % S_A[k], r[k], k))
    for k in S_B:
        out.append((HM, 'B%d' % S_B[k], r[k], k))
    out.append((HM, 'B%d' % S_M['cbola'], r['cbola'], 'cbola'))
    for i in range(NI):
        out.append((HE, 'D%d' % (E_I0 + i), r['neto'][i], 'neto%d' % i))
    for j in range(NP):
        for k in R:
            out.append((HE, '%s%d' % (PCOL[j], R[k]), r['p%d' % j][k], '%s[%d]' % (k, j)))
    return out


def _distinto(x, v):
    if isinstance(v, str):
        return x != v
    return not isinstance(x, (int, float)) or abs(float(x) - v) > 0.01


def autotest(destino):
    from pycel import ExcelCompiler
    err, n = [], 0
    wv = openpyxl.load_workbook(destino, data_only=True)
    wf = openpyxl.load_workbook(destino)
    base = valores_defecto()
    j2 = dict(base, bolas=120, peso=200, merma=0.05, h_l=1700, s_l=55, l_l=3, p_h=1.30, hid=0.62, sal_b=0.03, lev_b=0.0017,
              metodo=METODOS[1], tax=0.07, obj=0.28, tol=0.03, caja=0.9, pvp0=19.0, pvp6=25.0, p1=7.5, m1=0.05, g0_2=2.0, g7_5=3.5)
    j3 = dict(base, bolas=24, peso=300, merma=0.0, h_l=1500, s_l=60, p_s=0.9, p_l=4.0, hid=0.55, obj=0.2, tol=0.1, caja=0.6,
              pvp3=24.0, p2=20.0, m2=0.1, g1_1=2.5, g3_7=1.5)
    # 1. caché y pycel del juego por defecto; pycel de los otros dos con set_value
    xl = ExcelCompiler(filename=destino)
    r0 = calcular(base)
    for hoja, c, v, lab in celdas_salida(r0):
        n += 1
        for et, x in (('pycel', xl.evaluate("'%s'!%s" % (hoja, c))), ('caché', wv[hoja][c].value)):
            if _distinto(x, v):
                err.append('defecto · %s!%s (%s): %s %r ≠ esperado %r' % (hoja, c, lab, et, x, v))
    for nombre, jj in (('juego 2', j2), ('juego 3', j3)):
        r = calcular(jj)
        xl2 = ExcelCompiler(filename=destino)
        for hoja, c, v, lab in celdas_salida(r0):
            xl2.evaluate("'%s'!%s" % (hoja, c))
        for k, v in jj.items():
            if v != base[k]:
                h, c = CELDA[k]
                xl2.set_value("'%s'!%s" % (h, c), v)
        for hoja, c, v, lab in celdas_salida(r):
            n += 1
            x = xl2.evaluate("'%s'!%s" % (hoja, c))
            if _distinto(x, v):
                err.append('%s · %s!%s (%s): pycel %r ≠ esperado %r' % (nombre, hoja, c, lab, x, v))
    # 2. Σ ingredientes = masa (A y B) y B con los % de A reproduce A
    for nombre, jj in (('defecto', base), ('juego 2', j2), ('juego 3', j3)):
        r = calcular(jj)
        if abs(r['agua'] + r['harina'] + r['sal'] + r['lev'] - r['masa']) > 0.5:
            err.append('%s: Σ ingredientes del A ≠ masa' % nombre)
        if abs(r['harina_b'] + r['agua_b'] + r['sal_bg'] + r['lev_bg'] - r['masa']) > 0.5:
            err.append('%s: Σ ingredientes del B ≠ masa' % nombre)
        if abs(r['masa_lb'] * jj['g_lb'] - r['masa']) > 1e-6 or abs(r['harina_lb'] * jj['g_lb'] - r['harina']) > 1e-6:
            err.append('%s: el eco en lb no cuadra con los gramos' % nombre)
        rb = calcular(dict(jj, hid=r['ph_h'], sal_b=r['ph_s'], lev_b=r['ph_l']))
        for a, b in (('harina', 'harina_b'), ('agua', 'agua_b'), ('sal', 'sal_bg'), ('lev', 'lev_bg'), ('coste_a', 'coste_b'),
                     ('cbola_a', 'cbola_b')):
            n += 1
            if abs(rb[a] - rb[b]) > 0.01:
                err.append('%s: el B con los %% del A no reproduce el A (%s %.4f ≠ %s %.4f)' % (nombre, a, rb[a], b, rb[b]))
    xl3 = ExcelCompiler(filename=destino)                                       # lo mismo, con pycel sobre la hoja
    for k in S_B:
        xl3.evaluate("'%s'!B%d" % (HM, S_B[k]))
    for k, ck in (('hid', 'ph_h'), ('sal_b', 'ph_s'), ('lev_b', 'ph_l')):
        xl3.set_value("'%s'!B%d" % (HM, M_[k]), r0[ck])
    for a, b in (('harina', 'harina_b'), ('agua', 'agua_b'), ('sal', 'sal_bg'), ('lev', 'lev_bg'), ('cbola_a', 'cbola_b')):
        n += 1
        xa, xb = xl3.evaluate("'%s'!B%d" % (HM, S_A[a])), xl3.evaluate("'%s'!B%d" % (HM, S_B[b]))
        if abs(xa - xb) > 0.01:
            err.append('pycel: el B con los %% del A no reproduce el A (%s %.4f ≠ %s %.4f)' % (a, xa, b, xb))
    # 3. valores con sentido
    marg = r0['p0']
    if not 0.15 <= marg['fc'] <= 0.30:
        err.append('food cost de la Margherita %.1f %% fuera de 15-30 %%' % (marg['fc'] * 100))
    if r0['sem'] != SEM_OK:
        err.append('el valor por defecto no cae dentro de las proporciones tradicionales')
    if calcular(dict(base, h_l=2000))['sem'] != SEM_NO or calcular(dict(base, peso=300))['sem'] != SEM_NO:
        err.append('el semáforo no se pone rojo fuera de las proporciones')
    if base['tax'] != 0 or base['g_lb'] != 453.6 or base['oz_lb'] != 16:
        err.append('constantes por defecto: sales tax 0 %%, 453.6 g/lb, 16 oz/lb (%r %r %r)' % (base['tax'], base['g_lb'], base['oz_lb']))
    pizzas = list(PIZZAS)
    if any(re.search(r'authentic|verace|vpn', p, re.I) for p in pizzas):
        err.append('una pizza de la carta promete «authentic»/«verace»/«VPN»')
    esperado = OrderedDict([('Margherita', 18), ('Marinara', 15), ('Diavola', 20), ('Capricciosa', 21), ('Quattro Formaggi', 21),
                            ('Ham & Mushroom', 20), ('Bufalina', 23), ('Sausage & Broccoli Rabe', 22)])
    if OrderedDict((k, v[0]) for k, v in PIZZAS.items()) != esperado or base['caja'] != 0.75:
        err.append('carta o caja distintas de D145')
    # 4. formato, protección, DV, fórmulas
    if motor.gate_formatos(wf, FICHERO):
        err.append('gate_formatos: %r' % motor.gate_formatos(wf, FICHERO)[:3])
    dvs = dict((str(rng), dv) for dv in wf[HM].data_validations.dataValidation for rng in dv.sqref.ranges)
    if len(wf[HE].data_validations.dataValidation) < 3:
        err.append('faltan DV en %s' % HE)
    dv = dvs.get('B%d' % M_['bolas'])
    if dv is None or dv.type != 'whole' or str(dv.formula1) != '1':
        err.append('DV del número de bolas (entero > 0) ausente o incorrecta')
    dv = dvs.get('B%d' % M_['metodo'])
    if dv is None or dv.type != 'list':
        err.append('DV de lista del método ausente')
    for k in ('peso', 'h_l', 's_l', 'p_h'):
        if dvs.get('B%d' % M_[k]) is None:
            err.append('DV ausente en %s' % k)
    textos = []
    for w in wf.worksheets:
        if w.page_setup.paperSize != 1 or not w.sheet_properties.pageSetUpPr.fitToPage or not w.protection.sheet \
                or w.oddFooter.center.text != FOOTER:
            err.append('%s: Letter + ajuste + protección + pie' % w.title)
        for row in w.iter_rows():
            for c in row:
                v = c.value
                if v == '':
                    err.append('%s!%s: cadena vacía' % (w.title, c.coordinate))
                if isinstance(v, str) and c.data_type != 'f':
                    textos.append((w.title, c.coordinate, v))
                if isinstance(v, str) and c.data_type == 'f':
                    textos += [(w.title, c.coordinate + ' (fórmula)', lit) for lit in re.findall(r'"((?:[^"]|"")*)"', v)
                               if any(ch.isalpha() for ch in lit)]
                if c.data_type == 'f':
                    if wv[w.title][c.coordinate].value is None:
                        err.append('%s!%s: fórmula sin caché' % (w.title, c.coordinate))
                    sin_refs = re.sub(r'"[^"]*"', '', re.sub(r"'[^']+'!|\$?[A-Z]{1,2}\$?\d+", '', str(c.value)))
                    for num in re.findall(r'(?<![A-Za-z])\d+\.?\d*', sin_refs):
                        if num not in ('0', '1'):
                            err.append('%s!%s: número tecleado %s en %s' % (w.title, c.coordinate, num, c.value))
                if c.fill.fgColor.rgb in ('00E8F5E9', 'FFE8F5E9') and c.protection.locked:
                    err.append('%s!%s: verde bloqueada' % (w.title, c.coordinate))
                fmt = c.number_format or ''
                if '€' in fmt or '$' in fmt:
                    err.append('%s!%s: símbolo de moneda en el formato (D8)' % (w.title, c.coordinate))
        for dv_ in w.data_validations.dataValidation:
            for t in (dv_.error, dv_.errorTitle, dv_.prompt, dv_.promptTitle):
                if t:
                    textos.append((w.title, 'DV %s' % dv_.sqref, t))
    p = wf.properties
    for campo in ('title', 'subject', 'keywords', 'description', 'category'):
        textos.append(('docProps', campo, getattr(p, campo) or ''))
    for h, c, t in textos:
        r = [x for x in M.restos_espanol(t) if not x[0].startswith('oficio')]
        if r:
            err.append('%s!%s: restos de español %s' % (h, c, r[:1]))
        if M.no_latinos(t):
            err.append('%s!%s: caracteres no latinos' % (h, c))
        for rx in PROHIBIDOS:
            if re.search(rx, t, re.I):
                err.append('%s!%s: prohibido /%s/ en %r' % (h, c, rx, t[:60]))
        if re.search(r'(?i)\bvalue\b.*\$\d|\bsave \$|total value|\bwas \$|\d+\s?% off\b', t):
            err.append('%s!%s: ancla de precio' % (h, c))
    todo = '\n'.join(t for _h, _c, t in textos)
    for frase in ('does NOT make your pizza "VPN-certified"', 'no federal rule governs the term "Neapolitan-style"',
                  'TRADITIONAL NEAPOLITAN RATIOS (EU TSG spec, also used by the AVPN)', '453.6 g per lb', '16 oz per lb',
                  'Sales tax included in your menu prices (%)', SEM_OK, SEM_NO,
                  "Kit estimate (Oct 2026): replace with your distributor's price"):
        if frase not in todo:
            err.append('falta el texto %r' % frase)
    if M.FIRMA_EN not in todo or VERSION_LINEA not in todo or M.AVISO_EN not in todo:
        err.append('falta la firma, la línea de versión o el aviso')
    if wf[HE]['A5'].value != 'Sales tax included in your menu prices (%)' or wf[HE]['B5'].value != 0:
        err.append('Food Cost!B5 no es el sales tax = 0 %')
    if wf.sheetnames != [HM, HE, HI]:
        err.append('pestañas %r' % wf.sheetnames)
    if NI < 20:
        err.append('menos de 20 ingredientes')
    if p.creator != 'AI Chef Pro' or p.description != URL or 'AI Chef Pro' not in (p.keywords or '') \
            or not (p.title or '').startswith(TITULO):
        err.append('docProps D20: creator=%r title=%r description=%r keywords=%r' % (p.creator, p.title, p.description, p.keywords))
    for x in err[:80]:
        print('ERROR', x)
    print('autotest pizza dough & food cost calculator: %d comprobaciones (pycel + caché por defecto; pycel en 2 juegos más; B↔A) '
          'sobre %d textos · %s' % (n, len(textos), 'VERDE' if not err else 'ROJO (%d errores)' % len(err)))
    print('  dough by default: %d balls × %g g (%.1f oz) → %.0f g (%.1f lb) · %.2f l water · flour %.0f g (%.1f lb) · dough cost $%.2f '
          '· PER BALL $%.4f' % (base['bolas'], base['peso'], r0['peso_oz'], r0['masa'], r0['masa_lb'], r0['litros'], r0['harina'],
                               r0['harina_lb'], r0['coste_a'], r0['cbola_a']))
    for j, nom in enumerate(pizzas):
        q = r0['p%d' % j]
        print('  %-24s cost $%.3f · food cost %5.1f %% · %s' % (nom, q['c_tot'], q['fc'] * 100, q['sem']))
    return 1 if err else 0


if __name__ == '__main__':
    ap = argparse.ArgumentParser()
    ap.add_argument('--salida', required=True, help='output folder')
    ap.add_argument('--autotest', action='store_true')
    a = ap.parse_args()
    destino = os.path.join(a.salida, FICHERO)
    if not a.autotest and generar(destino):
        sys.exit('ERROR: fallos de pycel en inject_cache')
    sys.exit(autotest(destino))
