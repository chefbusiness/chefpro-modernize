#!/usr/bin/env python3
"""
semilla_hel.py — Plan de Negocio: Heladería Artesanal (ES 1.0; SPEC D4/D13/D14): prepara el ORIGEN que
`planes-v2_0/main.py` convierte en el plan financiero y el checklist de la heladería. Copia adaptada de
`parrillero-2_2/semilla_par.py` (F1-camino §3).

Copia los dos xlsx de la Panadería ES 2.2 publicados (`dl/plan-negocio-panaderia/`, SOLO LECTURA) con los nombres de
fichero de la heladería y escribe lo que el módulo de contenido no alcanza (F1-camino §1):
  · plan: valores de `0. Supuestos` (grupo_a conserva lo verde que trae el libro), títulos de hoja (Inversión A1/A2,
    PyG A2, Punto Equilibrio A2), las 19 partidas de inversión (filas 7-25: rótulo, importe, nota, ¿IVA?), los 10
    costes fijos propios del P&L (filas 26-35: rótulo e importe; la nota la pone el contenido) y vacía `PyG!H4:H11`
    (la línea partida «Pan común sobre la línea» de la panadería), y los docProps description/keywords;
  · checklist: título, pestañas, cabeceras de fase y las tareas de `datos_checklist.py` con el pie «Tareas
    completadas:» debajo; `Instrucciones!A3` con el pid; DV, formato condicional y contadores los rehace el motor.
Corre en el VPS (openpyxl). Después: `main.py --producto plan-negocio-heladeria --dry-run --origen <salida>`.

    /root/venv-guias/bin/python semilla_hel.py --salida /root/tmp-hel-f2/motor/origen

Sesión Claude Code, 25-oct-2026.
"""
import argparse
import copy
import importlib
import os
import shutil
import sys

import openpyxl
from openpyxl.styles import PatternFill

import datos_checklist as DC

AQUI = os.path.dirname(os.path.abspath(__file__))
REPO = os.path.abspath(os.path.join(AQUI, '..', '..', '..'))
PAN = os.path.join(REPO, 'astro-site', 'public', 'dl', 'plan-negocio-panaderia')
PLAN_PAN, CHECK_PAN = 'plan-financiero-panaderia.xlsx', 'checklist-apertura-panaderia.xlsx'
PLAN, CHECK = 'plan-financiero-heladeria.xlsx', 'checklist-apertura-heladeria.xlsx'
PID = 'plan-negocio-heladeria'

TITULOS_PLAN = {
    ('Inversión Inicial', 'A1'): 'PLAN DE NEGOCIO: HELADERÍA ARTESANAL',
    ('Inversión Inicial', 'A2'): 'Inversión inicial estimada — heladería artesanal con obrador propio',
    ('PyG 3 Años', 'A2'): 'Escenario realista — heladería artesanal',
    ('Punto Equilibrio', 'A2'): 'Punto de equilibrio — heladería artesanal',
}

#: Filas 7-25 de 'Inversión Inicial' (molde de la panadería): (rótulo, importe, nota, ¿lleva IVA?). Research §3: 19
#: partidas, 113.900 € (obra 32.500 · maquinaria 70.900 · no amortizable 10.500) + imprevistos del 10 % sobre la obra
#: (fila 28, por fórmula) = 117.150 €. Sin palabras de recuento en el rótulo (motor.RX_RECUENTO) ni «sala» (VOCABULARIO).
INVERSION = [
    ('Proyecto técnico y dirección de obra', 2500,
     'Proyecto del obrador y de la tienda, memoria de la actividad y certificado final', 'Sí'),
    ('Obra y acondicionamiento del obrador y la tienda', 18500,
     'Obrador alicatado, suelo antideslizante, aseo y falsos techos. Estimación del kit (reforma de 15.000-50.000 € '
     'en las guías publicadas)', 'Sí'),
    ('Instalación eléctrica trifásica, cuadro y boletín', 5000,
     'Unos 30 kW en trifásica para vitrinas, pasteurizador, mantecadora, abatidor y climatización. Pregunta la '
     'potencia disponible ANTES de firmar: una acometida nueva tarda', 'Sí'),
    ('Fontanería, desagües, climatización y ventilación', 4500,
     'Desagües del obrador, lavamanos de accionamiento no manual y salida de aire para los condensadores de las '
     'máquinas', 'Sí'),
    ('Rótulo y escaparate', 2000,
     '«Artesano» en el rótulo sólo con la calificación artesanal de tu comunidad; si no, «ELABORACIÓN PROPIA»', 'Sí'),
    ('Pasteurizador de 30 litros', 14600,
     'Pasteurizador de 15/30 l por ciclo, 4 kW: 14.601,75 € sin IVA en la tarifa de un distribuidor (Diamond '
     'SPS/30A, octubre de 2026). Pasteuriza la mezcla a 85 °C y la enfría a 4 °C', 'Sí'),
    ('Mantecadora horizontal de 6-12 litros por ciclo', 14000,
     'Estimación del kit en la parte baja de lo publicado (12.000-35.000 €), porque las marcas profesionales no '
     'publican tarifa: pide al menos dos presupuestos antes de cerrar el préstamo. Una vertical de entrada de 26-30 '
     'l/h cuesta unos 4.200 € sin IVA, si empiezas con poco', 'Sí'),
    ('Abatidor de temperatura para cubetas', 5500,
     'Lleva el helado recién mantecado a −18 °C en el corazón antes de pasarlo a la vitrina o al armario', 'Sí'),
    ('Armarios de conservación y de maduración', 5800,
     'Dos armarios a −20/−25 °C para el helado y uno a +2/+4 °C para madurar la mezcla y guardar la materia prima',
     'Sí'),
    ('Mesas inox, fregadero, básculas, refractómetro, sondas, cubetas y carapinas', 6500,
     'Mesas y estanterías inox, fregadero, lavamanos, turbo, básculas, refractómetro, sondas de temperatura, 60 '
     'cubetas de 5 l (44 € cada una) y 10 carapinas', 'Sí'),
    ('Vitrina ventilada de 18 cubetas', 10500,
     'Dos módulos ventilados de 10 cubetas, de −20/−9 °C y con gas R290 (unos 5.090 € cada uno en la tarifa del '
     'fabricante): 18 sabores y dos huecos para el sabor del día o la reserva. Con 18 sabores en verano y 10-12 en '
     'invierno basta en una ciudad media', 'Sí'),
    ('Mostrador, lavaporcionador, cafetera, crepera, gofrera y chocolatera', 6000,
     'Mostrador de venta, lavaporcionador con agua corriente, cafetera y molino, y el equipo del invierno. La '
     'crepera y la gofrera dependen de la salida de humos que permita tu ordenanza', 'Sí'),
    ('Mobiliario del local y de la terraza', 5000,
     'Mesas y sillas para 20 plazas dentro y 5 mesas de terraza en temporada', 'Sí'),
    ('TPV con Veri*factu, datáfono y báscula', 1800,
     'Software de facturación conforme a Veri*factu, datáfono y báscula para el helado al peso', 'Sí'),
    ('Menaje, uniformes y cartelería de alérgenos', 1200,
     'Copas, cucharas y espátulas (una por cubeta), uniformes y el cartel que dice dónde pedir la información de '
     'alérgenos (RD 126/2015)', 'Sí'),
    ('Formación técnica de heladería', 2000,
     'Curso de balanceo de mezclas: la diferencia entre un helado cremoso y uno arenoso o duro en la vitrina', 'Sí'),
    ('Marca, web, fotos y lanzamiento', 3500,
     'Identidad, web con la carta de sabores, sesión de fotos y campaña de apertura', 'Sí'),
    ('Constitución de la SL, licencia, tasas, registro sanitario y APPCC', 2000,
     'SL unipersonal, declaración responsable, tasas municipales, registro sanitario autonómico y redacción del '
     'plan APPCC. Las tasas no llevan IVA; los honorarios sí: cambia la columna de al lado si tu presupuesto es '
     'sobre todo de honorarios', 'No'),
    ('Stock inicial de materias primas y envases', 3000,
     'Leche, nata, azúcares, pastas, fruta, barquillos, tarrinas y envases para llevar', 'Sí'),
]
assert len(INVERSION) == 19 and sum(x[1] for x in INVERSION) == 113900, sum(x[1] for x in INVERSION)

#: Filas 26-35 de 'PyG 3 Años' (costes fijos propios). Research §3: 10 partidas, 10.250 €/año, una por fila del molde
#: (no se suprimen filas: desplazaría el P&L). Sin palabras de recuento en el rótulo. Las claves de
#: `contenido….FIJOS` son estos rótulos normalizados.
FIJOS = {
    26: ('Gestoría laboral, fiscal y contable', 1800),
    27: ('Mantenimiento del frío (contrato con frigorista certificado)', 1500),
    28: ('Marketing y publicidad de temporada', 2500),
    29: ('Software, teléfono e internet', 700),
    30: ('Limpieza, lavandería, EPI y análisis de autocontrol', 800),
    31: ('Tasa de terraza (pon la de tu ayuntamiento)', 1200),
    32: ('Gestión de residuos', 350),
    33: ('Desinsectación, desratización y desinfección (DDD)', 450),
    34: ('Derechos de autor por música ambiental (SGAE y AGEDI-AIE)', 500),
    35: ('Prevención de riesgos laborales y vigilancia de la salud', 450),
}
assert sum(x[1] for x in FIJOS.values()) == 10250


def _contenido():
    sys.path.insert(0, os.path.join(os.path.dirname(AQUI), 'planes-v2_0'))
    import grupo_a
    return grupo_a, importlib.import_module('contenido_plan_negocio_heladeria.a')


def _supuestos(wb):
    """Los valores de `0. Supuestos` del caso heladería: `grupo_a` CONSERVA lo que el libro ya trae en las celdas
    verdes y sólo usa el valor del contenido en una celda vacía; sin este paso el caso seguiría siendo el de la
    panadería (180 × 5,50 € × 310)."""
    grupo_a, contenido = _contenido()
    coords = dict((t[0], t[1]) for t in grupo_a.SUPUESTOS_BASE)
    ws = wb['0. Supuestos']
    escritas = 0
    for clave, t in contenido.SUPUESTOS.items():
        if t[2] is None:
            continue
        assert clave in coords, clave
        ws[coords[clave]] = t[2]
        escritas += 1
    return escritas


def _docprops(wb):
    """El motor no reescribe `description` ni `keywords` y heredaba los de la panadería."""
    wb.properties.description = 'aichef.pro/' + PID
    wb.properties.keywords = 'plan de negocio heladería artesanal, AI Chef Pro'


def plan(salida):
    destino = os.path.join(salida, PLAN)
    shutil.copy2(os.path.join(PAN, PLAN_PAN), destino)
    wb = openpyxl.load_workbook(destino)
    print('  supuestos escritos:', _supuestos(wb))
    _docprops(wb)
    for (hoja, coord), texto in TITULOS_PLAN.items():
        wb[hoja][coord] = texto
    ws = wb['Inversión Inicial']
    for i, (rot, imp, nota, iva) in enumerate(INVERSION):
        r = 7 + i
        assert not str(ws.cell(r, 1).value).lower().startswith(('fianza', 'alquiler', 'imprevistos')), r
        ws.cell(r, 1).value, ws.cell(r, 2).value, ws.cell(r, 4).value, ws.cell(r, 5).value = rot, imp, nota, iva
    assert str(ws['A26'].value).lower().startswith('fianza'), ws['A26'].value
    assert str(ws['A27'].value).lower().startswith('alquiler'), ws['A27'].value
    assert str(ws['A28'].value).lower().startswith('imprevistos'), ws['A28'].value
    ws = wb['PyG 3 Años']
    assert str(ws['A25'].value).startswith('Seguros') and str(ws['A36'].value).startswith('Amortización')
    for r, (rot, imp) in FIJOS.items():
        ws.cell(r, 1).value, ws.cell(r, 2).value = rot, imp
        ws.cell(r, 6).value = None          # la nota la pone el contenido (FIJOS)
    # la línea partida de la panadería («Pan común sobre la línea (%)» en H4 y su 0,76 en H10): sin línea partida
    # grupo_a no la reescribe ni consta que la borre (F1-camino §1)
    # IVA soportado de la compra de comida: en la panadería era una FÓRMULA (peso del pan común × 4 %); aquí es la
    # celda verde de `IVA_COMPRAS` (8,2 %). Se escribe el valor en el origen para que el motor no la cuente como una
    # fórmula perdida (control «el post-proceso PERDIÓ fórmulas que el fichero ya tenía»).
    _g, contenido = _contenido()
    assert str(ws['A13'].value).startswith('Coste de mercancía — comida'), ws['A13'].value
    assert str(ws['G13'].value).startswith('='), ws['G13'].value
    ws['G13'] = contenido.IVA_COMPRAS['comida']
    assert 'pan' in str(ws['H4'].value).lower(), ws['H4'].value
    for r in range(4, 12):
        c = ws.cell(r, 8)
        c.value = None
        c.fill = PatternFill(fill_type=None)
    wb.save(destino)
    return destino


def _copiar_estilo(origen, destino):
    destino.font, destino.border = copy.copy(origen.font), copy.copy(origen.border)
    destino.fill, destino.alignment = copy.copy(origen.fill), copy.copy(origen.alignment)
    destino.number_format, destino.protection = origen.number_format, copy.copy(origen.protection)


def checklist(salida):
    destino = os.path.join(salida, CHECK)
    shutil.copy2(os.path.join(PAN, CHECK_PAN), destino)
    wb = openpyxl.load_workbook(destino)
    fases = [ws for ws in wb.worksheets if ws.title[:2] in ('F1', 'F2', 'F3', 'F4', 'F5', 'F6')]
    assert len(fases) == 6
    for ws, (titulo, cabecera, tareas) in zip(fases, DC.FASES):
        ws.title = titulo
        ws['A1'], ws['A2'] = DC.TITULO, cabecera
        assert ws['B4'].value and 'TRÁMITE' in ws['B4'].value.upper(), ws['B4'].value
        modelo = [copy.copy(ws.cell(5, c)) for c in range(1, 6)]
        alto = ws.row_dimensions[5].height
        pie = [copy.copy(ws.cell(r, c)) for r in range(5, ws.max_row + 1) for c in range(1, 6)
               if isinstance(ws.cell(r, c).value, str) and 'completad' in ws.cell(r, c).value.lower()]
        assert pie, ws.title
        fila_pie = pie[0].row
        estilo_pie = [copy.copy(ws.cell(fila_pie, c)) for c in range(1, 6)]
        # vaciar la zona de tareas y el pie (valores, DV y formato condicional los rehace el motor)
        for r in range(5, ws.max_row + 1):
            for c in range(1, 6):
                ws.cell(r, c).value = None
            ws.row_dimensions[r].height = None
        ws.data_validations.dataValidation = []
        ws.conditional_formatting = type(ws.conditional_formatting)()
        for i, (tarea, resp, plazo, nota) in enumerate(tareas):
            r = 5 + i
            for c, v in enumerate(('☐', tarea, resp, plazo, nota), start=1):
                cel = ws.cell(r, c)
                _copiar_estilo(modelo[c - 1], cel)
                cel.value = v
            ws.row_dimensions[r].height = alto
        r = 5 + len(tareas) + 1
        for c in range(1, 6):
            _copiar_estilo(estilo_pie[c - 1], ws.cell(r, c))
        ws.cell(r, 2).value = 'Tareas completadas:'
        ws.cell(r, 3).value = '=COUNTIF(A5:A%d,"✓")' % (4 + len(tareas))
        ws.cell(r, 4).value = 'de'
        ws.cell(r, 5).value = '=COUNTIF(B5:B%d,"?*")-COUNTIF(A5:A%d,"N/A")' % (4 + len(tareas), 4 + len(tareas))
    # `motor.asegurar_instrucciones` no reescribe la hoja si ya existe: la de la panadería nombraba su producto
    ins = wb['Instrucciones']
    assert 'plan-negocio-panaderia' in str(ins['A3'].value), ins['A3'].value
    ins['A3'] = 'Este fichero forma parte del producto «%s» de AI Chef Pro.' % PID
    _docprops(wb)
    wb.save(destino)
    return destino


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--salida', required=True)
    a = ap.parse_args()
    if os.path.isdir(a.salida):
        shutil.rmtree(a.salida)
    os.makedirs(a.salida)
    print('plan      →', plan(a.salida))
    print('checklist →', checklist(a.salida), '(%d tareas: %s)' % (
        sum(len(t) for _, _, t in DC.FASES), '/'.join(str(len(t)) for _, _, t in DC.FASES)))


if __name__ == '__main__':
    main()
