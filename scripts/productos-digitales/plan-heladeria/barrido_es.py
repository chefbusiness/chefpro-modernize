#!/usr/bin/env python3
"""barrido_es.py — Heladería Artesanal ES 1.0 (SPEC D12; copia de `parrillero-2_2/barrido_es.py`): barre los xlsx/docx que se
le pasen buscando restos del molde FT (camión, ventanilla, generador, aparcamiento, food truck), restos de la panadería
(semilla: pan, bollería, harina, masa madre, hornada, levadura, madrugada, transacción…) y los prohibidos de la SPEC
(España 2026, validado, freelance, carné de manipulador, tarifa plana, CNAE, RGSEAA para el minorista, licencia de
apertura, versión anterior…): la lista es `tandas_docx.PROHIBIDOS`. Lista celda/párrafo y texto; en las fórmulas sólo
mira los literales entre comillas. Corre en el VPS. Uso: barrido_es.py <fichero>… · exit 1 si hay coincidencias.
Sesión Claude Code, 25-oct-2026."""
import os
import re
import sys

import docx
import openpyxl

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import tandas_docx as T                                                            # noqa: E402

RX = re.compile('|'.join(T.PROHIBIDOS), re.I)
n = 0
for p in [x for x in sys.argv[1:] if not x.startswith('--')]:
    if p.endswith('.xlsx'):
        wb = openpyxl.load_workbook(p)
        it = [(ws.title + '!' + c.coordinate, c.value) for ws in wb.worksheets for row in ws.iter_rows() for c in row
              if isinstance(c.value, str)]
        pr = wb.properties
        it += [('docProps.' + k, getattr(pr, k)) for k in ('title', 'subject', 'description', 'keywords', 'creator')
               if isinstance(getattr(pr, k), str)]
    else:
        d = docx.Document(p)
        it = [('p%d' % i, x.text) for i, x in enumerate(d.paragraphs)]
        it += [('tabla%d' % i, c.text) for i, t in enumerate(d.tables) for row in t.rows for c in row.cells]
        cp = d.core_properties
        it += [('docProps.' + k, getattr(cp, k)) for k in ('title', 'subject', 'comments', 'keywords', 'author')
               if isinstance(getattr(cp, k), str)]
    for loc, t in it:
        txt = ' '.join(re.findall(r'"([^"]*)"', t)) if t.startswith('=') else t
        m = RX.search(txt)
        if m:
            n += 1
            print('%s · %s · «%s» · …%s…' % (p.split('/')[-1][:28], loc, m.group(0), txt[max(0, m.start() - 50):m.end() + 50]))
print('barrido: %d coincidencias' % n)
sys.exit(1 if n else 0)
