#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
aplicar_en.py — Restaurant Checklists Kit Pro (EN) · montaje de los 11 xlsx desde el Kit de Tareas Recurrentes v2.0
publicado (`dl/kit-tareas/`, SOLO LECTURA).

SPEC que manda: `SPEC.md` (§1 D1-D23, §2 mapas, §3 reglas US/UK, §7 paso 3). Sesión Claude Code, 17-oct-2026.
Copia adaptada de `staff-kit/aplicar_en.py`: misma maquinaria (renombrado de hojas con reescritura de referencias,
literales de fórmula/CF/DV por CLAVES, traducción POR APARICIÓN, US Letter + pie, docProps, comillas rectas, huella
de idempotencia, `inject_cache` AL FINAL, `--dry-run`, `--idempotencia`). Aquí no se redacta nada: los textos salen
de `textos_en/G*.json` (subagentes Anthropic, regla 1bis) y de `mapas.py`.

    python3 aplicar_en.py                      # escribe en astro-site/public/dl/restaurant-checklist-templates/
    python3 aplicar_en.py --dry-run [--salida <carpeta>] [--mes October] [--idempotencia] [--json informe.json]
    python3 aplicar_en.py --dry-run --parcial  # lo que aún no tenga traducción lleva la marca «PENDING-EN cNNNN»
    python3 aplicar_en.py --dry-run --parcial --textos <carpeta vacía>   # ensayo solo con la marca de relleno

Orden por libro, siempre desde una copia limpia del ES publicado (→ idempotente):
   1. carga del xlsx ES y renombrado de hojas (mapas.HOJAS, §2.2) reescribiendo la única referencia de fórmula a
      otra hoja (09 «Cash Drawer Closing»!C48 → 'Cash Drawer Opening'!C23); áreas y títulos de impresión salen con
      el nombre nuevo (openpyxl los regenera; «Valentine's Day» → 'Valentine''s Day'!);
   2. fórmulas: las 15 de mapas.FORMULAS_POR_CELDA (D12, multiplicadores en dólares); el resto, HOJAS + CLAVES
      (CLAVES = {N/A}: un literal con letras sin clave aborta; «?*», «✓», «>0» no cambian);
   3. textos POR APARICIÓN (textos_es.json `donde`): traducir → textos_en; mapa-glosario → GLOSARIO; mapa-hoja →
      HOJAS; mapa-clave → CLAVES; horas de 24 h de celda → mapas.hora_us (D9); línea de versión y de marca (D19); pie;
   4. DV: listas por CLAVES, numéricas tal cual, mensajes por aparición; CF: literales por CLAVES;
   5. capas regeneradas: título interior (D5, Instructions!B2) y POR_CELDA (caja D12, calendario D14, servicios
      contratados D11, parásitos D10, fichajes);
   6. comillas rectas; ajuste de texto y alto de fila donde el EN ya no cabe (sin tocar anchos, nunca baja un alto);
   7. formatos (D13: «#,##0.00 €» → «#,##0.00»; formatos propios sin uso, renombrados); US Letter en las 47 (D17);
   8. docProps (D19), título interior comprobado;
   9. cobertura: toda aparición «regenerar» del ES tiene que haberla escrito una capa (si no, aborta);
  10. guardado con el nombre EN (mapas.FICHEROS) e inject_cache.py AL FINAL sobre los 11 (nada guarda después).

La protección de las 36 hojas de trabajo (sin contraseña), DV, CF, merges, paneles y anchos viajan intactos.

Térmica: todo en SERIE; `istats` antes de cada libro y de cada inject_cache (espera si ≥ 62 °C hasta < 60 °C). En el
VPS no hay istats y la guarda no espera. En el VPS `/dev/null` está roto: nada se redirige allí.
"""
import argparse
import copy
import datetime
import glob
import json
import logging
import math
import os
import re
import shutil
import subprocess
import sys
import tempfile
import time
from collections import OrderedDict

AQUI = os.path.dirname(os.path.abspath(__file__))
SCRIPTS = os.path.dirname(AQUI)
REPO = os.path.dirname(os.path.dirname(SCRIPTS))
sys.path.insert(0, AQUI)

import openpyxl                                              # noqa: E402
from openpyxl.utils import get_column_letter                 # noqa: E402

import mapas                                                 # noqa: E402
import extraer_textos                                        # noqa: E402

logging.disable(logging.CRITICAL)

ORIGEN = extraer_textos.ORIGEN                                                 # dl/kit-tareas (solo lectura)
DESTINO_REAL = os.path.join(REPO, 'astro-site', 'public', 'dl', mapas.SLUG)   # restaurant-checklist-templates
INJECT = os.path.join(SCRIPTS, 'inject_cache.py')
TEXTOS_EN = os.path.join(AQUI, 'textos_en')

VERSION = mapas.VERSION_ES                                            # 2.0 (contenido v2.0 del ES)
SUBJECT = mapas.DOCPROPS_EN['subject']                                # D19
KEYWORDS = mapas.DOCPROPS_EN['keywords']
DESCRIPTION = mapas.DOCPROPS_EN['description']
CATEGORY = mapas.DOCPROPS_EN['category']
CREATOR = 'AI Chef Pro'
PIE_EN = mapas.FOOTER_EN                                              # «AI Chef Pro · aichef.pro · Page &P of &N»
MESES = ['January', 'February', 'March', 'April', 'May', 'June', 'July', 'August',
         'September', 'October', 'November', 'December']
VERDE = extraer_textos.VERDE
FORMATOS = OrderedDict(mapas.FORMATOS_EN)                             # D13
INVARIABLES = set(extraer_textos.INVARIABLES)
RX_VERSION_ES = extraer_textos.RX_VERSION
HOJA_INSTR_EN = mapas.HOJAS['Instrucciones']
CLAVES = mapas.CLAVES
CAMPO_DV = {'dv-error-titulo': 'errorTitle', 'dv-error': 'error', 'dv-prompt-titulo': 'promptTitle', 'dv-prompt': 'prompt'}
DV_MENSAJES_EN = getattr(mapas, 'DV_MENSAJES_EN', OrderedDict())      # arreglos por aparición (si los hubiera)
RELLENO = 'PENDING-EN %s'                                             # --parcial: marca de lo que aún no está
RX_RELLENO = re.compile(r'PENDING-EN c\d{4}')
PT_LINEA = 1.36                                                       # alto de línea ≈ 1.36 × cuerpo (pt)
ALTO_DEFECTO = 15.0
ALTO_MAX = 409.0


class Aborta(Exception):
    pass


# ==========================================================================
# Térmica (regla del Mac: se apaga por encima de ~65 °C; en el VPS no hay istats)
# ==========================================================================
def termica(umbral=62, reanuda=60):
    def leer():
        try:
            out = subprocess.run(['istats', 'cpu', 'temp', '--value-only'], stdout=subprocess.PIPE,
                                 stderr=subprocess.PIPE, universal_newlines=True, timeout=20).stdout
            return float(re.findall(r'[\d.]+', out)[0])
        except Exception:
            return None
    t = leer()
    if t is None or t < umbral:
        return t
    while t is not None and t >= reanuda:
        print('  [térmica] CPU %.1f °C: espero 30 s' % t, flush=True)
        time.sleep(30)
        t = leer()
    return t


# ==========================================================================
# Datos
# ==========================================================================
def cargar_traducciones(carpeta, parcial):
    """textos_en/G<n>.json = {id: en}. En --parcial, un fichero a medio escribir (JSON inválido) se ignora con aviso."""
    tx = OrderedDict()
    avisos = []
    for p in sorted(glob.glob(os.path.join(carpeta, 'G[0-9].json'))):
        try:
            d = json.load(open(p, encoding='utf-8'))
        except ValueError as e:
            if not parcial:
                raise Aborta('%s no es JSON válido: %s' % (os.path.basename(p), e))
            avisos.append('%s ilegible (¿a medio escribir?): ignorado' % os.path.basename(p))
            continue
        for k, v in d.items():
            if not isinstance(v, str) or not v.strip():
                raise Aborta('%s: traducción vacía o no textual para %s' % (os.path.basename(p), k))
            if k in tx and tx[k] != v:
                raise Aborta('traducción contradictoria para %s en %s' % (k, os.path.basename(p)))
            tx[k] = v
    return tx, avisos


def cargar_datos(parcial=False, textos=TEXTOS_EN):
    te = json.load(open(os.path.join(AQUI, 'textos_es.json'), encoding='utf-8'))
    censo = json.load(open(os.path.join(AQUI, 'censo_es.json'), encoding='utf-8'))
    tx, avisos = cargar_traducciones(textos, parcial)
    por_id = {c['id']: c for c in te['cadenas']}
    ajenas = sorted(set(tx) - set(por_id))
    if ajenas:
        raise Aborta('textos_en trae ids que no están en textos_es.json: %r' % ajenas[:5])
    gm = sorted(k for k in tx if por_id[k]['grupo'] == 'GM')
    if gm:
        raise Aborta('textos_en traduce cadenas GM (las escribe mapas.py): %r' % gm[:5])
    faltan = [c['id'] for c in te['cadenas'] if c['grupo'] != 'GM' and c['id'] not in tx]
    if faltan and not parcial:
        raise Aborta('%d cadenas G1-G2 sin traducción (p. ej. %r)' % (len(faltan), faltan[:5]))
    celdas, dvmsg = {}, {}
    for c in te['cadenas']:
        for d in c['donde']:
            if d['t'] == 'celda':
                celdas[(d['f'], d['h'], d['c'])] = (c['id'], d['tr'], c['es'])
            elif d['t'] in CAMPO_DV:
                dvmsg[(d['f'], d['h'], d['c'], CAMPO_DV[d['t']])] = (c['id'], d['tr'], c['es'])
    return {'te': te, 'censo': censo, 'tx': tx, 'celdas': celdas, 'dvmsg': dvmsg, 'faltan': faltan,
            'parcial': parcial, 'avisos': avisos, 'por_id': por_id}


def en_traducido(datos, ident):
    """EN de una cadena G1-G2; en --parcial, la marca de relleno si aún no está (None si falta y no es parcial)."""
    en = datos['tx'].get(ident)
    if en is None and datos['parcial']:
        return RELLENO % ident
    return en


def mes_por_defecto():
    return MESES[datetime.date.today().month - 1]


def linea_version(mes):
    """D19: «Version 2.0 · <mes> 2026 · aichef.pro/en/digital-products/restaurant-checklist-templates · info@…»."""
    return 'Version %s · %s %d · %s · info@aichef.pro' % (VERSION, mes, datetime.date.today().year, DESCRIPTION)


def regenerado_esperado(corto, h_es, coord, es, mes):
    """EN que una capa escribe en una aparición «regenerar» de celda (lo usan la cobertura y G4)."""
    k = (corto, h_es, coord)
    if k in mapas.POR_CELDA:
        return mapas.POR_CELDA[k]
    if h_es == 'Instrucciones' and coord == 'B2':
        return mapas.titulo_interior(next(f for f in mapas.LIBROS_ES if extraer_textos.corto(f) == corto))
    if RX_VERSION_ES.match(es):
        return linea_version(mes)
    if es == mapas.MARCA_ES:
        return mapas.MARCA_EN
    if mapas.RX_HORA_CELDA.match(es):
        return mapas.hora_us(es)
    return None


# ==========================================================================
# Fórmulas: hojas y literales
# ==========================================================================
RX_STR = re.compile(r'"(?:[^"]|"")*"')
RX_REF = re.compile(r"(?:'((?:[^']|'')+)'|([^\W\d][\w.]*))!")


def ref_hoja(nombre):
    """Nombre de hoja tal como va en una fórmula (entre comillas si hace falta; el apóstrofo se dobla)."""
    if re.fullmatch(r'[A-Za-z_][A-Za-z0-9_.]*', nombre):
        return nombre
    return "'" + nombre.replace("'", "''") + "'"


def clave_en(lit, donde):
    """Literal ES → EN por CLAVES: estricto con los que llevan letras o «€»; los demás no cambian."""
    if not mapas.necesita_clave(lit):
        return lit
    if lit not in CLAVES:
        raise Aborta('literal sin clave EN %r en %r' % (lit, donde[:90]))
    return CLAVES[lit]


class Transformador:
    """Reescribe una fórmula ES a EN: referencias de hoja (fuera de literales) y literales (dentro)."""

    def __init__(self, mapa_hojas):
        self.mapa = mapa_hojas
        self.en = set(mapa_hojas.values())
        self.hojas_citadas = 0
        self.literales = 0

    def refs(self, s):
        def rep(m):
            nombre = (m.group(1) if m.group(1) is not None else m.group(2)).replace("''", "'")
            if nombre in self.mapa:
                self.hojas_citadas += 1
                return ref_hoja(self.mapa[nombre]) + '!'
            if nombre in self.en:
                return m.group(0)
            raise Aborta('hoja desconocida en fórmula: %r en %r' % (nombre, s[:80]))
        return RX_REF.sub(rep, s)

    def formula(self, f):
        out, pos = [], 0
        for m in RX_STR.finditer(f):
            out.append(self.refs(f[pos:m.start()]))
            lit = m.group(0)[1:-1].replace('""', '"')
            nuevo = clave_en(lit, f)
            if nuevo != lit:
                self.literales += 1
            out.append('"' + nuevo.replace('"', '""') + '"')
            pos = m.end()
        out.append(self.refs(f[pos:]))
        return ''.join(out)

    def celda(self, v):
        return '=' + self.formula(v[1:]) if v.startswith('=') else self.formula(v)


def formula_en(v, corto, h_es, coord, T):
    """Fórmula EN de una celda: FORMULAS_POR_CELDA (D12) o HOJAS + CLAVES. Devuelve (nueva, por_celda)."""
    k = (corto, h_es, coord)
    if k in mapas.FORMULAS_POR_CELDA:
        es, en = mapas.FORMULAS_POR_CELDA[k]
        if v != es:
            raise Aborta('%s %s!%s: FORMULAS_POR_CELDA espera %r y el ES trae %r' % (corto, h_es, coord, es, v))
        return en, True
    return T.celda(v), False


def rango(ref):
    c1, r1, c2, r2 = openpyxl.utils.cell.range_boundaries(ref)
    for r in range(r1, r2 + 1):
        for c in range(c1, c2 + 1):
            yield '%s%d' % (get_column_letter(c), r)


def rango_lista(sqref):
    out = []
    for rg in str(sqref).split():
        out.extend(rango(rg))
    return out


def renombrar(wb):
    mapa = OrderedDict((ws.title, mapas.HOJAS[ws.title]) for ws in wb.worksheets)
    for ws in wb.worksheets:
        ws.title = mapa[ws.title]
    if wb.sheetnames != list(mapa.values()):
        raise Aborta('renombrado: %r ≠ %r' % (wb.sheetnames, list(mapa.values())))
    return mapa


def reescribir_formulas(L, T):
    n = pc = 0
    for ws in L.wb.worksheets:
        h_es = L.inv[ws.title]
        for c in list(ws._cells.values()):
            if c.data_type == 'f' and isinstance(c.value, str):
                nuevo, es_pc = formula_en(c.value, L.corto, h_es, c.coordinate, T)
                pc += es_pc
                if nuevo != c.value:
                    c.value = nuevo
                    n += 1
    usadas = sum(1 for k in mapas.FORMULAS_POR_CELDA if k[0] == L.corto)
    if pc != usadas:
        raise Aborta('%s: FORMULAS_POR_CELDA aplicadas %d de %d' % (L.corto, pc, usadas))
    L.rep['formulas_reescritas'] = n
    L.rep['formulas_por_celda'] = pc


# ==========================================================================
# Textos por aparición
# ==========================================================================
class Libro:
    """Estado de un libro en construcción: celdas escritas por capas (cobertura de «regenerar»)."""

    def __init__(self, wb, corto, mapa, datos, rep):
        self.wb, self.corto, self.mapa, self.datos, self.rep = wb, corto, mapa, datos, rep
        self.inv = {en: es for es, en in mapa.items()}
        self.escritas = set()                 # (hoja_es, celda)
        self.faltan = rep.setdefault('faltan', [])

    def ws(self, hoja_es):
        return self.wb[self.mapa[hoja_es]]

    def poner(self, hoja_es, coord, valor, es_esperado=None):
        c = self.ws(hoja_es)[coord]
        if c.data_type == 'f' or (isinstance(c.value, str) and c.value.startswith('=')):
            raise Aborta('%s %s!%s es una fórmula (%r): una capa no la pisa' % (self.corto, hoja_es, coord, c.value))
        if es_esperado is not None and not es_esperado(c.value):
            raise Aborta('%s %s!%s: el ES trae %r, no lo esperado' % (self.corto, hoja_es, coord, c.value))
        c.value = valor
        self.escritas.add((hoja_es, coord))
        return c

    def en_de(self, ident, tr, es, lugar):
        if tr in ('traducir', 'localizar'):
            en = en_traducido(self.datos, ident)
            if en is None:
                self.faltan.append('%s %s %r' % (lugar, ident, es[:60]))
                return es
            if RX_RELLENO.match(en):
                self.rep['relleno'] = self.rep.get('relleno', 0) + 1
            return en
        if tr == 'mapa-clave':
            return CLAVES[es]
        if tr == 'mapa-hoja':
            return mapas.HOJAS[es]
        if tr == 'mapa-glosario':
            return mapas.GLOSARIO[es]
        return None                                              # regenerar: lo escribe una capa


def es_texto(v):
    return isinstance(v, str) and bool(v)


def traducir_textos(L, mes):
    n = OrderedDict([('textos_traducidos', 0), ('glosario', 0), ('horas_us', 0), ('versiones', 0), ('marcas', 0),
                     ('pies', 0)])
    regenerar = []
    for ws in L.wb.worksheets:
        h_es = L.inv[ws.title]
        for c in list(ws._cells.values()):
            v = c.value
            if not isinstance(v, str) or c.data_type == 'f':
                continue
            if RX_VERSION_ES.match(v):
                L.poner(h_es, c.coordinate, linea_version(mes))
                n['versiones'] += 1
                continue
            if v == mapas.MARCA_ES:
                L.poner(h_es, c.coordinate, mapas.MARCA_EN)
                n['marcas'] += 1
                continue
            if mapas.RX_HORA_CELDA.match(v):                     # D9: «07:00» → «7:00 AM» (texto)
                L.poner(h_es, c.coordinate, mapas.hora_us(v))
                n['horas_us'] += 1
                continue
            k = (L.corto, h_es, c.coordinate)
            if k not in L.datos['celdas']:
                if any(ch.isalpha() for ch in v) and v not in INVARIABLES:
                    raise Aborta('%s %s!%s: texto %r fuera de textos_es.json' % (L.corto, h_es, c.coordinate, v[:60]))
                continue                                         # sin letras / invariable (✓, —, #, N/A…)
            ident, tr, es = L.datos['celdas'][k]
            if es != v:
                raise Aborta('%s %s!%s: textos_es.json dice %r y el ES trae %r' % (L.corto, h_es, c.coordinate,
                                                                                es[:40], v[:40]))
            en = L.en_de(ident, tr, es, '%s %s!%s' % (L.corto, ws.title, c.coordinate))
            if en is None:
                regenerar.append((h_es, c.coordinate))
                continue
            if tr == 'mapa-glosario':
                n['glosario'] += 1
            if en != v:
                c.value = en
                n['textos_traducidos'] += 1
        for parte in ('oddHeader', 'oddFooter', 'evenHeader', 'evenFooter', 'firstHeader', 'firstFooter'):
            hf = getattr(ws, parte)
            for pos in ('left', 'center', 'right'):
                t = getattr(hf, pos).text
                if not t:
                    continue
                if t != mapas.FOOTER_ES:
                    raise Aborta('%s %s %s.%s inesperado: %r' % (L.corto, ws.title, parte, pos, t))
                getattr(hf, pos).text = PIE_EN
                n['pies'] += 1
    L.rep.update(n)
    return regenerar


# ==========================================================================
# DV y CF
# ==========================================================================
def dv_censo(L, h_es, sq):
    for d in L.datos['censo']['libros'][L.rep['es']]['hojas_detalle'][h_es]['dv']:
        if d['sqref'] == sq:
            return d
    raise Aborta('%s %s: DV %s no está en el censo' % (L.corto, h_es, sq))


def dv_lista(items):
    return '"' + ','.join(items) + '"'


def dv_lista_en(items, donde):
    """Ítems de una lista literal → EN por CLAVES («✓,—,N/A»: ✓ y — sin letras; N/A igual en inglés)."""
    out = []
    for x in items:
        if x == '' or not mapas.necesita_clave(x):
            out.append(x)
            continue
        if x not in CLAVES:
            raise Aborta('%s: ítem de DV sin clave EN %r' % (donde, x))
        out.append(CLAVES[x])
    return dv_lista(out)


def reescribir_dv_cf(L, T):
    n = OrderedDict([('dv_listas', 0), ('dv_personalizadas', 0), ('dv_ref', 0), ('dv_mensajes', 0),
                     ('dv_mensajes_fijados', 0), ('cf_reglas', 0), ('cf_text', 0)])
    usados = set()
    for ws in L.wb.worksheets:
        h_es = L.inv[ws.title]
        for dv in ws.data_validations.dataValidation:
            sq = str(dv.sqref)
            f1 = dv.formula1 or ''
            dv_censo(L, h_es, sq)
            if dv.type == 'list' and f1.startswith('"'):
                dv.formula1 = dv_lista_en(f1.strip('"').split(','), '%s %s DV %s' % (L.corto, h_es, sq))
                if len(dv.formula1) > 257:
                    raise Aborta('%s %s: DV %s > 255 caracteres' % (L.corto, h_es, sq))
                n['dv_listas'] += 1
            elif dv.type == 'custom' and f1:
                dv.formula1 = T.formula(f1)
                n['dv_personalizadas'] += 1
            elif f1:
                dv.formula1 = T.refs(f1)
                if dv.formula2:
                    dv.formula2 = T.refs(dv.formula2)
                n['dv_ref'] += 1
            for attr in ('error', 'errorTitle', 'prompt', 'promptTitle'):
                v = getattr(dv, attr)
                if not v:
                    continue
                k = (L.corto, h_es, sq, attr)
                if k not in L.datos['dvmsg']:
                    raise Aborta('%s %s DV %s %s fuera de textos_es.json: %r' % (L.corto, h_es, sq, attr, v[:50]))
                ident, tr, es = L.datos['dvmsg'][k]
                if es != v:
                    raise Aborta('%s %s DV %s %s: textos_es.json dice %r' % (L.corto, h_es, sq, attr, es[:40]))
                en = L.en_de(ident, tr, es, '%s DV %s %s' % (ws.title, sq, attr))
                if en is None:
                    raise Aborta('%s %s DV %s %s marcado «regenerar»' % (L.corto, h_es, sq, attr))
                fijo = DV_MENSAJES_EN.get((L.corto, h_es, sq, attr))
                if fijo is not None:
                    en = fijo
                    usados.add((L.corto, h_es, sq, attr))
                    n['dv_mensajes_fijados'] += 1
                setattr(dv, attr, en)
                n['dv_mensajes'] += 1
        for cf in ws.conditional_formatting:
            for regla in cf.rules:
                if regla.formula:
                    regla.formula = [T.formula(x) for x in regla.formula]
                    n['cf_reglas'] += 1
                if regla.text:
                    regla.text = clave_en(regla.text, 'CF %s %s' % (ws.title, cf.sqref))
                    n['cf_text'] += 1
    sin_usar = [k for k in DV_MENSAJES_EN if k[0] == L.corto and k not in usados]
    if sin_usar:
        raise Aborta('%s: DV_MENSAJES_EN sin DV donde aplicarse: %r' % (L.corto, sin_usar))
    L.rep.update(n)


# ==========================================================================
# Capas regeneradas (mapas.py)
# ==========================================================================
def capa_titulo(L, fname_es):
    L.poner('Instrucciones', 'B2', mapas.titulo_interior(fname_es), es_texto)


def capa_por_celda(L):
    n = 0
    for (c, h, coord), v in mapas.POR_CELDA.items():
        if c == L.corto:
            L.poner(h, coord, v, es_texto)
            n += 1
    L.rep['por_celda'] = n


COMILLAS = {'\u2019': "'", '\u2018': "'", '\u201c': '"', '\u201d': '"'}


def normalizar_comillas(L):
    n = 0

    def norm(v):
        for a, b in COMILLAS.items():
            v = v.replace(a, b)
        return v
    for ws in L.wb.worksheets:
        for c in list(ws._cells.values()):
            if isinstance(c.value, str) and c.data_type != 'f' and any(a in c.value for a in COMILLAS):
                c.value = norm(c.value)
                n += 1
        for dv in ws.data_validations.dataValidation:
            for attr in ('error', 'errorTitle', 'prompt', 'promptTitle'):
                v = getattr(dv, attr)
                if v and any(a in v for a in COMILLAS):
                    setattr(dv, attr, norm(v))
                    n += 1
    L.rep['comillas_normalizadas'] = n


# ==========================================================================
# Ajuste de texto y alto de fila (riesgo F1: textos EN más largos en columnas fijas). Sin tocar anchos.
# Estimación conservadora (la de staff-kit): un carácter por unidad de ancho de columna a 11 pt.
# Solo actúa en celdas cuyo texto EN difiere del ES (la maqueta ES, aprobada, manda en lo demás):
#   · sin ajuste y el EN no cabe en su caja (combinada, o propia + vecinas vacías por donde desborda) → ajuste ON;
#   · con ajuste en una fila de alto fijo, o celda combinada de una sola fila (Excel no autoajusta combinadas):
#     si el EN pide más líneas que el ES, alto = líneas EN × 1.36 × cuerpo + el aire del ES (≤ 6 pt); nunca baja.
#     Las filas de alto automático sin combinar las autoajusta Excel al abrir (como en el ES).
# ==========================================================================
def ancho_col(ws, idx):
    for d in ws.column_dimensions.values():
        if d.min and d.max and d.min <= idx <= d.max and d.width:
            return d.width
    return ws.sheet_format.defaultColWidth or 8.43


def mapa_merges(ws):
    anclas, cubiertas = {}, set()
    for m in ws.merged_cells.ranges:
        anclas[(m.min_row, m.min_col)] = (m.min_col, m.min_row, m.max_col, m.max_row)
        for r in range(m.min_row, m.max_row + 1):
            for c in range(m.min_col, m.max_col + 1):
                if (r, c) != (m.min_row, m.min_col):
                    cubiertas.add((r, c))
    return anclas, cubiertas


def _pt(cel):
    return cel.font.sz if cel.font is not None and cel.font.sz else 11.0


def ancho_texto(t, cel):
    """Ancho (unidades de columna) de la línea más larga: 1 carácter = 1 unidad a 11 pt (+8 % en negrita)."""
    k = _pt(cel) / 11.0 * (1.08 if cel.font is not None and cel.font.b else 1.0)
    return max(len(seg) for seg in t.split('\n')) * k


def n_lineas(t, ancho, cel):
    k = _pt(cel) / 11.0 * (1.08 if cel.font is not None and cel.font.b else 1.0)
    cpl = max(1, int(ancho / k))
    return sum(max(1, -(-len(seg) // cpl)) for seg in t.split('\n'))


def caja(ws, cel, anclas, cubiertas):
    """(ancho de la caja, filas, combinada, ancho disponible sin ajuste) o None si la celda está cubierta."""
    pos = (cel.row, cel.column)
    if pos in cubiertas:
        return None
    if pos in anclas:
        c1, r1, c2, r2 = anclas[pos]
        w = sum(ancho_col(ws, c) for c in range(c1, c2 + 1))
        return w, r2 - r1 + 1, True, w
    w = ancho_col(ws, cel.column)
    disp = w
    ultima = max(ws.max_column, 1)
    for c in range(cel.column + 1, ultima + 1):
        vecina = ws._cells.get((cel.row, c))
        if (vecina is not None and vecina.value is not None) or (cel.row, c) in cubiertas or (cel.row, c) in anclas:
            break
        disp += ancho_col(ws, c)
    return w, 1, False, disp


def palabra_mas_larga(t, cel):
    k = _pt(cel) / 11.0 * (1.08 if cel.font is not None and cel.font.b else 1.0)
    return max([len(w) for w in t.split()] or [0]) * k


def objetivo_ajuste(ws_en, ws_es, cel, anclas, cubiertas):
    """Lo que la celda EN necesita: (activar ajuste, alto mínimo o None, desborde sin arreglo). Misma cuenta en
    aplicar_en (capa_ajuste) y en G3. La maqueta ES manda: solo se actúa donde el ES cabía y el EN ya no."""
    t = cel.value
    cel_es = ws_es._cells.get((cel.row, cel.column))             # sin crear celdas en el ES
    es = cel_es.value if cel_es is not None else None
    if not isinstance(t, str) or cel.data_type == 'f' or t == es or cel_es is None:
        return False, None, False
    cj = caja(ws_en, cel, anclas, cubiertas)
    if cj is None:
        return False, None, False
    ancho, filas, combinada, disponible = cj
    wrap_es = bool(cel_es.alignment and cel_es.alignment.wrap_text)
    necesita = False
    if not wrap_es and ancho_texto(t, cel) > disponible + 1e-9:
        disp_es = caja(ws_es, cel_es, anclas, cubiertas)[3]
        if not (isinstance(es, str) and ancho_texto(es, cel_es) <= disp_es + 1e-9):
            return False, None, False                            # el ES tampoco cabía: maqueta heredada
        if palabra_mas_larga(t, cel) > ancho + 1e-9 or filas > 1:
            return False, None, True                             # ni con ajuste: hay que acortar el EN
        necesita = True
    if not (wrap_es or necesita) or filas > 1:
        return necesita, None, False
    dim_es = ws_es.row_dimensions.get(cel.row)
    alto_es = dim_es.height if dim_es is not None else None
    if alto_es is None and not combinada:
        return necesita, None, False                             # fila automática sin combinar: la ajusta Excel
    l_en = n_lineas(t, ancho, cel)
    l_es = n_lineas(es, ancho, cel_es) if (isinstance(es, str) and wrap_es) else 1
    if alto_es is None:                                          # combinada de alto automático: Excel no la ajusta
        alto = l_en * _pt(cel) * PT_LINEA
        return necesita, (min(ALTO_MAX, math.ceil(alto * 4) / 4.0) if alto > ALTO_DEFECTO else None), False
    if l_en <= l_es:
        return necesita, None, False
    linea = _pt(cel) * PT_LINEA                                  # líneas EN + el aire del ES (hasta 6 pt)
    alto = l_en * linea + max(0.0, min(6.0, alto_es - l_es * linea))
    return necesita, min(ALTO_MAX, math.ceil(max(alto, alto_es) * 4) / 4.0), False


def capa_ajuste(L, wes):
    n = OrderedDict([('ajuste_activado', 0), ('altos_subidos', 0)])
    desbordes = L.rep.setdefault('desbordes', [])
    for ws in L.wb.worksheets:
        ws_es = wes[L.inv[ws.title]]
        anclas, cubiertas = mapa_merges(ws)
        for cel in list(ws._cells.values()):
            necesita, alto, desborde = objetivo_ajuste(ws, ws_es, cel, anclas, cubiertas)
            if desborde:
                desbordes.append('%s!%s %r' % (ws.title, cel.coordinate, cel.value[:50]))
            if necesita:
                al = copy.copy(cel.alignment)
                al.wrap_text = True
                cel.alignment = al
                n['ajuste_activado'] += 1
            if alto is not None:
                actual = ws.row_dimensions[cel.row].height
                if alto > (actual or ALTO_DEFECTO):
                    ws.row_dimensions[cel.row].height = alto
                    n['altos_subidos'] += 1
    L.rep.update(n)


# ==========================================================================
# Formatos y papel (D13, D17)
# ==========================================================================
def formatos_y_papel(L):
    wb = L.wb
    n = OrderedDict([('fmt_' + k, 0) for k in FORMATOS])
    for ws in wb.worksheets:
        for c in list(ws._cells.values()):
            f = c.number_format
            if f in FORMATOS:
                c.number_format = FORMATOS[f]
                n['fmt_' + f] += 1
        ws.page_setup.paperSize = ws.PAPERSIZE_LETTER
    # formatos propios que ya no usa ninguna celda (los del ES) → variantes sin «€»
    nf = wb._number_formats
    usados = set(nf)
    for i, f in enumerate(list(nf)):
        if '€' in f:
            nuevo = FORMATOS.get(f, '#,##0.00')
            while nuevo in usados:
                nuevo += ';@'
            usados.add(nuevo)
            list.__setitem__(nf, i, nuevo)
    if len(set(nf)) != len(nf):
        raise Aborta('formatos propios repetidos: %r' % list(nf))
    nf._dict = {v: i for i, v in enumerate(nf)}
    nf.clean = True
    L.rep.update(n)


def metadatos(L, fname_en):
    p = L.wb.properties
    p.title = mapas.TITULOS[fname_en]                 # D5
    p.subject = SUBJECT
    p.keywords = KEYWORDS
    p.description = DESCRIPTION
    p.category = CATEGORY
    p.creator = CREATOR
    p.lastModifiedBy = CREATOR
    L.rep['docprops'] = OrderedDict((k, getattr(p, k)) for k in
                                    ('title', 'subject', 'creator', 'keywords', 'description', 'category'))


def comprobar_titulo(L, fname_es):
    """D5: el título interior (Instructions!B2) es EXACTAMENTE mapas.titulo_interior. Nunca se parchea."""
    esperado = mapas.titulo_interior(fname_es)
    v = L.wb[HOJA_INSTR_EN]['B2'].value
    if v != esperado:
        raise Aborta('%s: Instructions!B2 %r ≠ D5 %r' % (fname_es, v, esperado))
    L.rep['titulo'] = v


# ==========================================================================
# Un libro
# ==========================================================================
def construir_libro(fname_es, carpeta, datos, mes):
    fname_en = mapas.FICHEROS[fname_es]
    corto = extraer_textos.corto(fname_es)
    rep = OrderedDict([('es', fname_es), ('en', fname_en)])
    wb = openpyxl.load_workbook(os.path.join(ORIGEN, fname_es))
    wes = openpyxl.load_workbook(os.path.join(ORIGEN, fname_es))                # referencia ES (solo lectura)
    mapa = renombrar(wb)                                                       # 1
    L = Libro(wb, corto, mapa, datos, rep)
    T = Transformador(mapa)
    reescribir_formulas(L, T)                                                  # 2
    regenerar = traducir_textos(L, mes)                                        # 3
    reescribir_dv_cf(L, T)                                                     # 4
    capa_titulo(L, fname_es)                                                   # 5
    capa_por_celda(L)
    rep['refs_hoja'] = T.hojas_citadas
    rep['literales'] = T.literales
    normalizar_comillas(L)                                                     # 6
    capa_ajuste(L, wes)
    formatos_y_papel(L)                                                        # 7
    metadatos(L, fname_en)                                                     # 8
    comprobar_titulo(L, fname_es)
    sin_capa = [x for x in regenerar if x not in L.escritas]                   # 9
    if sin_capa:
        raise Aborta('%s: apariciones «regenerar» que ninguna capa escribió: %r' % (fname_es, sin_capa[:8]))
    for h_es, coord in regenerar:
        es = wes[h_es][coord].value
        esp = regenerado_esperado(corto, h_es, coord, es, mes)
        if esp is None or L.ws(h_es)[coord].value != esp:
            raise Aborta('%s %s!%s: aparición «regenerar» con %r (esperado %r)' % (corto, h_es, coord,
                                                                                L.ws(h_es)[coord].value, esp))
    rep['regeneradas'] = len(regenerar)
    if rep['faltan'] and not datos['parcial']:
        raise Aborta('%s: %d textos sin traducción:\n  %s' % (fname_es, len(rep['faltan']), '\n  '.join(rep['faltan'][:20])))
    wb.save(os.path.join(carpeta, fname_en))                                  # 10
    return rep


def construir(carpeta, datos, mes, log=print):
    os.makedirs(carpeta, exist_ok=True)
    informe = []
    for fname_es in mapas.LIBROS_ES:
        termica()
        rep = construir_libro(fname_es, carpeta, datos, mes)
        informe.append(rep)
        log('  %-48s fx %3d (D12 %2d) · refs %d · textos %4d · glos %3d · horas %3d · por celda %3d · regen %3d'
            ' · ajuste %d/%d · desbordes %d · relleno %d'
            % (rep['en'], rep['formulas_reescritas'], rep['formulas_por_celda'], rep['refs_hoja'],
               rep['textos_traducidos'], rep['glosario'], rep['horas_us'], rep['por_celda'], rep['regeneradas'],
               rep['ajuste_activado'], rep['altos_subidos'], len(rep['desbordes']), rep.get('relleno', 0)))
        for d in rep['desbordes'][:5]:
            log('      desborde (acortar el EN): ' + d)
    return informe


# ==========================================================================
# Idempotencia (huella comparable, sin fechas de guardado)
# ==========================================================================
def digest(path):
    wb = openpyxl.load_workbook(path)
    p = wb.properties
    fuera = {'_props': (p.title, p.subject, p.keywords, p.description, p.category, p.creator)}
    for ws in wb.worksheets:
        celdas = {}
        for c in ws._cells.values():
            if c.value is None and not (c.fill is not None and c.fill.fill_type == 'solid'):
                continue
            relleno = str(c.fill.fgColor.rgb) if c.fill is not None and c.fill.fill_type == 'solid' else None
            celdas[c.coordinate] = (repr(c.value), c.number_format, relleno, bool(c.protection.locked),
                                    bool(c.alignment.wrap_text) if c.alignment else None)
        fuera[ws.title] = {
            'celdas': celdas,
            'merges': sorted(str(m) for m in ws.merged_cells.ranges),
            'dv': sorted('%s:%s:%s:%s:%s:%s:%s' % (dv.type, dv.formula1, dv.formula2, dv.sqref, dv.error,
                                                   dv.errorTitle, dv.prompt)
                         for dv in ws.data_validations.dataValidation),
            'cf': sorted('%s:%s:%s' % (cf.sqref, [r.formula for r in cf.rules], [r.text for r in cf.rules])
                         for cf in ws.conditional_formatting),
            'prot': (bool(ws.protection.sheet), ws.protection.password), 'area': str(ws.print_area),
            'titulos': str(ws.print_titles), 'papel': ws.page_setup.paperSize, 'pie': ws.oddFooter.center.text,
            'altos': sorted((r, d.height) for r, d in ws.row_dimensions.items() if d.height),
        }
    return fuera


def diferencias(a, b, fichero):
    out = []
    if a['_props'] != b['_props']:
        out.append('%s: propiedades' % fichero)
    for hoja in sorted((set(a) | set(b)) - {'_props'}):
        if hoja not in a or hoja not in b:
            out.append('%s:%s: hoja solo en una pasada' % (fichero, hoja))
            continue
        for k in ('merges', 'dv', 'cf', 'prot', 'area', 'titulos', 'papel', 'pie', 'altos'):
            if a[hoja][k] != b[hoja][k]:
                out.append('%s:%s: cambia %s' % (fichero, hoja, k))
        ca, cb = a[hoja]['celdas'], b[hoja]['celdas']
        for coord in sorted(set(ca) | set(cb)):
            if ca.get(coord) != cb.get(coord):
                out.append('%s:%s!%s: %s → %s' % (fichero, hoja, coord, ca.get(coord), cb.get(coord)))
    return out


# ==========================================================================
# inject_cache y verificación de la caché
# ==========================================================================
def inject_cache(carpeta, log=print):
    fuera = []
    for fname_es in mapas.LIBROS_ES:
        termica()
        p = os.path.join(carpeta, mapas.FICHEROS[fname_es])
        r = subprocess.run([sys.executable, '-W', 'ignore', INJECT, p], stdout=subprocess.PIPE,
                           stderr=subprocess.PIPE, universal_newlines=True)
        linea = r.stdout.strip().splitlines()[-1] if r.stdout.strip() else ''
        log('    ' + linea)
        fuera.append({'fichero': mapas.FICHEROS[fname_es], 'salida': linea, 'exit': r.returncode,
                      'stderr': r.stderr[-400:] if r.returncode else ''})
    return fuera


def verificar_cache(carpeta):
    out = OrderedDict()
    for fname_es in mapas.LIBROS_ES:
        p = os.path.join(carpeta, mapas.FICHEROS[fname_es])
        wf, wv = openpyxl.load_workbook(p), openpyxl.load_workbook(p, data_only=True)
        n = con_valor = 0
        for ws in wf.worksheets:
            for c in ws._cells.values():
                if c.data_type == 'f':
                    n += 1
                    if wv[ws.title][c.coordinate].value is not None:
                        con_valor += 1
        out[mapas.FICHEROS[fname_es]] = (n, con_valor)
    return out


def carpeta_dryrun(arg):
    if arg:
        return os.path.abspath(arg)
    s = os.environ.get('CLAUDE_SCRATCHPAD')
    if s:
        return os.path.join(os.path.abspath(s), 'checklists-dryrun')
    return os.path.join(REPO, '.work', mapas.SLUG, 'checklists-dryrun')


def main():
    ap = argparse.ArgumentParser(description='Restaurant Checklists Kit Pro (EN): montaje de los 11 xlsx')
    ap.add_argument('--dry-run', action='store_true', help='escribe en el scratchpad (o --salida), no en dl/')
    ap.add_argument('--salida', default=None, help='carpeta del dry-run')
    ap.add_argument('--mes', default=None, help='mes de la línea de versión D19 (por defecto, el actual)')
    ap.add_argument('--json', default=None, help='informe JSON')
    ap.add_argument('--idempotencia', action='store_true', help='2.ª construcción en un clon y 0 diferencias')
    ap.add_argument('--parcial', action='store_true',
                    help='(solo --dry-run) lo que no tenga traducción lleva la marca «PENDING-EN cNNNN»')
    ap.add_argument('--textos', default=TEXTOS_EN, help='carpeta de G<n>.json (por defecto textos_en/)')
    ap.add_argument('--sin-cache', action='store_true', help='(depuración) no corre inject_cache')
    args = ap.parse_args()
    if args.parcial and not args.dry_run:
        raise SystemExit('--parcial solo vale con --dry-run')

    mes = args.mes or mes_por_defecto()
    if mes not in MESES:
        raise SystemExit('--mes debe ser un mes en inglés: %s' % ', '.join(MESES))
    carpeta = carpeta_dryrun(args.salida) if args.dry_run else DESTINO_REAL
    if os.path.abspath(carpeta) == os.path.abspath(ORIGEN):
        raise SystemExit('ABORTADO: la salida no puede ser la carpeta ES')

    t0 = time.time()
    try:
        datos = cargar_datos(args.parcial, args.textos)
        for a in datos['avisos']:
            print('  aviso: ' + a)
        print('== 1/3 · montaje de los 11 libros → %s (mes %s)%s' % (
            carpeta, mes, ' · PARCIAL: %d cadenas G1-G2 con la marca de relleno' % len(datos['faltan'])
            if args.parcial else ''), flush=True)
        informe = construir(carpeta, datos, mes)
    except Aborta as e:
        raise SystemExit('ABORTADO: %s' % e)
    sobran = sorted(set(os.listdir(carpeta)) - set(mapas.FICHEROS.values()))
    if sobran:
        print('  aviso: ficheros ajenos en la salida: %s' % sobran)

    idem = {'ejecutada': False}
    if args.idempotencia:
        print('== 2/3 · idempotencia: 2.ª construcción en un clon', flush=True)
        clon = tempfile.mkdtemp(prefix='checklists-idem-')
        construir(clon, datos, mes, log=lambda *_: None)
        difs = []
        for n in mapas.FICHEROS.values():
            difs += diferencias(digest(os.path.join(carpeta, n)), digest(os.path.join(clon, n)), n)
        shutil.rmtree(clon)
        idem = {'ejecutada': True, 'diferencias': len(difs), 'detalle': difs[:30]}
        print('  diferencias 1.ª vs 2.ª construcción: %d' % len(difs), flush=True)
        for d in difs[:10]:
            print('    ' + d)

    print('== 3/3 · inject_cache (AL FINAL, sobre los 11)', flush=True)
    cache = [] if args.sin_cache else inject_cache(carpeta)
    ver = verificar_cache(carpeta) if cache else {}
    for f, (n, v) in ver.items():
        print('    data_only %-48s %4d fórmulas · %4d con valor' % (f, n, v))

    fallos = []
    if idem.get('diferencias'):
        fallos.append('idempotencia: %d diferencias' % idem['diferencias'])
    fallos += ['inject_cache %s: exit %s %s' % (c['fichero'], c['exit'], c['stderr'][-200:]) for c in cache if c['exit']]
    fallos += ['inject_cache %s: %s' % (c['fichero'], c['salida']) for c in cache
               if c['exit'] == 0 and 'fallos_pycel=0' not in c['salida']]
    salida = OrderedDict([('producto', mapas.SLUG), ('version', VERSION),
                          ('fecha', datetime.datetime.now().isoformat(timespec='seconds')),
                          ('modo', 'dry-run' if args.dry_run else 'real'), ('parcial', args.parcial),
                          ('carpeta', carpeta), ('mes', mes), ('libros', informe), ('idempotencia', idem),
                          ('inject_cache', cache), ('cache_data_only', ver), ('fallos', fallos),
                          ('segundos', round(time.time() - t0, 1))])
    if args.json:
        with open(args.json, 'w', encoding='utf-8') as fh:
            json.dump(salida, fh, ensure_ascii=False, indent=1, default=str)
        print('informe → %s' % args.json)
    if fallos:
        print('\nFALLOS:\n  ' + '\n  '.join(fallos))
        sys.exit(1)
    print('\nOK: 11 xlsx en %s (%.0f s)%s' % (carpeta, time.time() - t0,
                                             ' · PARCIAL (marca de relleno)' if args.parcial else ''))


if __name__ == '__main__':
    main()
