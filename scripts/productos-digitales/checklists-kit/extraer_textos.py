#!/usr/bin/env python3
"""
extraer_textos.py — Restaurant Checklists Kit Pro (EN) · cimientos de la F2-EN.

Copia adaptada de `staff-kit/extraer_textos.py` (sesión Claude Code, 17-oct-2026). Lee los 11 xlsx PUBLICADOS
del Kit de Tareas Recurrentes v2.0 (`astro-site/public/dl/kit-tareas/`, SOLO LECTURA; el docx de equivalencias
ES↔LatAm no se replica) y escribe, junto a este script (o en --salida):

  · censo_es.json   por libro y hoja, lo que contarán los gates de la F2-EN: fórmulas (huella sha1, PATRONES con
                    su recuento, texto por celda), referencias a hojas, literales en fórmulas, CF y DV, DV, CF,
                    gráficos, imágenes, merges, protección, paneles, áreas y títulos de impresión, papel, pies,
                    formatos (con €), textos por celda, tareas (filas numeradas con texto), celdas verdes y
                    desbloqueadas, docProps; y `dependencias`: lo que depende de España, por categoría.
  · textos_es.json  cadenas de texto únicas con su contexto, TRATAMIENTO, PISTAS, ADAPTAR (instrucción de
                    adaptación por celda), citas de pestañas/rótulos/ficheros y GRUPO de traducción
                    (G1-G2 para los dos subagentes; GM = no se traduce).

    python3 extraer_textos.py [--kit base|bar] [--origen DIR] [--salida DIR] [--md]

--kit (kit.py): sin él, el kit base de siempre (dl/kit-tareas → junto a este script); con --kit <v>, el vertical
(dl/<mapas.ORIGEN_ES> → <v>/, con <v>/mapas_<v>.py).

No traduce nada (regla 1bis: los textos EN los escriben subagentes Anthropic, nunca bridge.py).
Térmica: un solo proceso, en serie, sin builds ni navegador (11 libros, < 15 s).

Tratamientos (la cadena toma el más «fuerte»: traducir > localizar > mapa-hoja > mapa-clave > mapa-glosario >
regenerar):
  traducir       texto libre → subagente (G1-G2), con pistas y, si la celda lo pide, ADAPTAR.
  localizar      dato de ejemplo → subagente (este kit no tiene: ZONAS_EJEMPLO vacía).
  mapa-hoja      nombre de pestaña → `mapas.HOJAS`.
  mapa-clave     literal de fórmula/CF/DV → `mapas.CLAVES`.
  mapa-glosario  celda corta repetida (cabecera, zona, puesto, cuándo) → `mapas.GLOSARIO`.
  regenerar      lo escribe `aplicar_en.py`: título interior, versión, línea de marca, pie, docProps, horas
                 (`mapas.hora_us`), rótulos fijados (`mapas.POR_CELDA`: caja, calendario, servicios contratados).
"""
import argparse
import datetime
import hashlib
import json
import os
import re
import subprocess
import sys
import zipfile
from collections import OrderedDict, defaultdict

import openpyxl
from openpyxl.styles.numbers import is_date_format
from openpyxl.utils import get_column_letter, range_boundaries

AQUI = os.path.dirname(os.path.abspath(__file__))
REPO = os.path.abspath(os.path.join(AQUI, '..', '..', '..'))
sys.path.insert(0, AQUI)
import kit                                                       # noqa: E402
from kit import mapas                                            # noqa: E402
ORIGEN = kit.ORIGEN                                              # base: dl/kit-tareas · vertical: dl/<ORIGEN_ES>

VERDE = 'E8F5E9'
INVARIABLES = {'N/A', 'AI Chef Pro', '© 2026 AI Chef Pro · aichef.pro', 'Mise en Place'} | set(
    getattr(mapas, 'INVARIABLES_EXTRA', ()))
RX_VERSION = re.compile(r'^Versi[óo]n \d+\.\d+ · ')
RX_LITERAL = re.compile(r'"((?:[^"]|"")*)"')
RX_REF_HOJA = re.compile(r"(?:'((?:[^']|'')+)'|([^\W\d][\w.]*))!(?=\$?[A-Za-z]{1,3}\$?\d|\$?[A-Za-z]{1,3}:)")
RX_CITA = re.compile(r"«([^»]{1,60})»|'([^']{1,60})'|\"([^\"]{1,60})\"")
RX_FICHERO = re.compile(r'(?:BONUS-)?\d{2}-[a-z-]+\.xlsx')
RX_COMPARA = re.compile(r'[<>]=?\s*-?\d')
RX_HORA_FMT = re.compile(r'h{1,2}:mm', re.I)
ZONAS_EJEMPLO = {}

# Lo que depende de España (F1-inventario §1). id → (etiqueta, regex sobre el TEXTO, destino). «hora24» se mide
# sobre la celda entera (mapas.RX_HORA_CELDA); «puestos» por pertenencia a mapas.PUESTOS_ES.
DEPENDENCIAS = OrderedDict([
    ('celsius', ('Temperaturas en °C', r'°C', 'D10: °F + FDA Food Code (ADAPTAR o POR_CELDA)')),
    ('hora24', ('Horas de 24 h como celda («07:00»)', None, 'D9: `mapas.hora_us` (regenerar)')),
    ('hora_texto', ('Horas de 24 h dentro de una frase', r'\b\d{1,2}:\d{2}\b', 'D9: el traductor, reloj de 12 h')),
    ('moneda', ('Euros (€) en textos', r'€|\bEUR\b|euros?\b', 'D12-D13: dólares en la tabla de caja; sin símbolo en '
                                                            'el resto')),
    ('caja', ('Caja española: TPV, datáfono, Bizum, arqueo, descuadre', r'TPV|[Dd]atáfono|Bizum|Visa|Mastercard|'
              r'\bMC\b|[Aa]rqueo|X de apertura|[Dd]escuadre|DESCUADRE|[Ff]ondo de caja|[Ff]ondo [Aa]pertura',
              'D12: POS, card terminal, Z/X report, starting bank, over/short, propinas fuera')),
    ('appcc', ('APPCC y «Pack APPCC»', r'APPCC', 'D10: HACCP · «HACCP Food Safety Kit Pro»')),
    ('norma', ('Normativa y administración española', r'\bDDD\b|\bBIE\b|Verifactu|antifraude|[Ll]egionela|'
               r'[Gg]estor autorizado|[Pp]unto limpio|[Rr]etimbrado|homologada|habilitada|Ayuntamiento|'
               r'[Rr]egistros? (diario )?de jornada|4 años|ANISAKIS|[Aa]nisakis|[Ll]ejía|[Ff]rigorista|'
               r'[Mm]antenedor|Empresa DDD|Instalador gas|[Ll]aboratorio|póliza',
               'D10-D11: FDA Food Code, OSHA, NFPA 96, EPA 608, FLSA, normas locales (POR_CELDA / ADAPTAR)')),
    ('festivos', ('Festivos y costumbres españolas', r'Reyes|Roscón|Carnaval|Día del Padre|Semana Santa|vigilia|'
                  r'torrijas|bacalao|potaje|Comuniones|bautizos|Día del Trabajo|Día de la Madre|San Juan|Asunción|'
                  r'Todos los Santos|Constitución|Inmaculada|Nochebuena|Nochevieja|NOCHEVIEJA|campanadas|\buvas\b|'
                  r'\bcava\b|turrón|cotillón|setas|castañas|menú del día|[Vv]uelta al cole',
                  'D14: calendario US (`mapas.CALENDARIO`) y ADAPTAR en 06 «Christmas & NYE»')),
    ('temporada', ('Terraza y temporada alta', r'[Tt]erraza|TERRAZA|[Tt]emporada alta',
                   'D15: patio / patio season (Memorial Day–Labor Day)')),
    ('puestos', ('Puestos y zonas de la hostelería española (celda entera)', None,
                 'D8: `mapas.GLOSARIO` (line cook, prep cook, server, porter, MOD…)')),
    ('plataformas', ('Plataformas españolas', r'TripAdvisor|Google Maps', 'D16: Google, Yelp, TripAdvisor; Google '
                                                                         'Business Profile')),
    ('alergenos', ('Alérgenos (UE 14)', r'[Aa]lérgenos|ALÉRGENOS', 'D10: 9 alérgenos US; nota UK 14')),
    ('a4', ('Papel A4 citado', r'\bA4\b', 'D17: US Letter')),
    ('variante', ('Sufijo de variante «Restaurante Casual»', r'Restaurante Casual', 'D2: la EN no lleva sufijo')),
    ('ficheros', ('Citas a ficheros ES (.xlsx)', RX_FICHERO.pattern, 'D5: nombre de fichero EN (`mapas.FICHEROS`)')),
    ('excel_es', ('Menús de Excel en español', r'Revisar →|[Cc]lic(k)? derecho|Desproteger|Mover o copiar',
                  'Excel EN: «Review → Unprotect Sheet», «right-click → Insert», «Move or Copy»')),
])
DEPENDENCIAS.update(getattr(mapas, 'DEPENDENCIAS_EXTRA', {}))   # vertical: añade o sustituye (kit.py)

PISTAS = [
    (r'°C|[Tt]emperatura|cámara|congela', 'D10: °F primero (41 °F frío, 0 °F congelador, 135 °F caliente; FDA Food '
                                          'Code 3-501.16); coletilla «— log reading: ____ °F»'),
    (r'APPCC', 'D10: APPCC → HACCP; «Pack APPCC» → «HACCP Food Safety Kit Pro»'),
    (r'TPV|[Dd]atáfono|Z del TPV|X de apertura|[Aa]rqueo|[Dd]escuadre|[Ff]ondo de caja',
     'D12: TPV → POS; datáfono → card terminal; Z del TPV → POS Z report; X → X report; arqueo → cash count; '
     'descuadre → over/short; fondo de caja → starting bank'),
    (r'€|euros?', 'D13: sin símbolo de moneda («in your currency»)'),
    (r'[Cc]amarer|[Cc]ociner|[Jj]efe|[Jj]efa|[Aa]uxiliar|[Ee]ncargad|[Bb]arman|[Pp]artida|[Pp]astel',
     'D8 puestos US (mapas.GLOSARIO): camarero → server; cocinero → line cook; chef de partida → lead line cook; '
     'auxiliar → prep cook; jefe de cocina → head chef; jefe de sala / maître → FOH manager; encargado → shift '
     'manager; barman → bartender; partida → station; comanda → ticket; pase → the pass'),
    (r'Hora Límite|✓ Completada|\bFirma\b|\bNotas\b|Responsable,', 'D8: nombres de columna = mapas.GLOSARIO '
                                                                  '(Assigned to, Due by, ✓ Done, Initials, Notes)'),
    (r'[Ss]ala\b|SALA|[Cc]omedor', 'Glosario: sala → front of house (FOH) / dining room'),
    (r'[Tt]erraza|TERRAZA', 'D15: terraza → patio'),
    (r'[Cc]ubiertos?|[Cc]omensales|[Tt]icket medio', 'Glosario: cubiertos/comensales → covers/guests; ticket medio '
                                                     '→ average check'),
    (r'[Cc]ámara|CÁMARA', 'Glosario: cámara → walk-in cooler / walk-in freezer'),
    (r'[Pp]lancha|[Bb]año maría|[Ss]alamandra|[Cc]ampana', 'Glosario: plancha → flat-top griddle; baño maría → '
                                                          'steam table; salamandra → salamander; campana → hood'),
    (r'[Mm]ermas?', 'Glosario: mermas → waste (waste log)'),
    (r'[Aa]lérgenos|ALÉRGENOS', 'D10: los 9 alérgenos principales de EE. UU. (nota UK: 14)'),
    (r'[Rr]eseñas|TripAdvisor|Google Maps', 'D16: reseñas en Google, Yelp y TripAdvisor; Google Business Profile'),
    (r'[Ss]eñal|[Aa]nticipo', 'Glosario: señal / anticipo → deposit'),
    (r'[Bb]riefing', 'Glosario: briefing de servicio → pre-shift meeting (line-up)'),
    (r'Turno: ☐', 'Turnos US: «Shift: ☐ AM ☐ PM ☐ Overnight»; en el briefing «☐ Lunch ☐ Dinner»'),
    (r'Responsable( de turno| general)?:', 'Formularios: «Responsable de turno:» → «Shift manager:»; «Responsable:» '
                                           '→ «Manager:»'),
    (r'Restaurante Casual', 'D2: la EN no lleva el sufijo de variante: quita «— Restaurante Casual»'),
    (RX_FICHERO.pattern, 'D5: cada fichero citado → su nombre EN tal cual (ver cita_ficheros)'),
    (r'[«»]', 'D6: pestañas, rótulos y valores citados → su EN de cita_rotulos, entre comillas dobles rectas'),
    (r'Revisar →|[Cc]lic(k)? derecho|Desproteger|Mover o copiar', 'Excel EN: «Review → Unprotect Sheet», '
                                                                 '«right-click → Insert», «right-click → Move or '
                                                                 'Copy»'),
    (r'\bA4\b', 'D17: US Letter'),
    (r'\d,\d', 'Formato US: punto decimal y coma de miles'),
    (r'\b\d{1,2}:\d{2}\b', 'D9: horas en reloj de 12 h («6:00 PM», «11:30 PM», medianoche = «midnight»)'),
    (r'Kit de Tareas', 'D2: «Restaurant Checklists Kit Pro»'),
    (r'86', '«86» / «86s» es jerga US: se queda'),
]
# vertical (kit.py): PISTAS_CAMBIO sustituye el texto de una pista por su regex; PISTAS_EXTRA se añade al final
PISTAS = [(rx, getattr(mapas, 'PISTAS_CAMBIO', {}).get(rx, t)) for rx, t in PISTAS
          if rx not in getattr(mapas, 'PISTAS_FUERA', ())] + list(getattr(mapas, 'PISTAS_EXTRA', []))


# --------------------------------------------------------------------------
# utilidades
# --------------------------------------------------------------------------
def sha256(path):
    h = hashlib.sha256()
    with open(path, 'rb') as fh:
        for b in iter(lambda: fh.read(65536), b''):
            h.update(b)
    return h.hexdigest()


def corto(fichero):
    m = re.match(r'^(BONUS-)?(\d+)', fichero)
    return ('B' if m.group(1) else '') + m.group(2)


def tiene_letra(s):
    return any(ch.isalpha() for ch in s)


def literales(formula):
    return [m.group(1).replace('""', '"') for m in RX_LITERAL.finditer(formula or '')]


def refs_hoja(formula):
    sin_lit = RX_LITERAL.sub('""', formula or '')
    return [(m.group(1) or m.group(2)).replace("''", "'") for m in RX_REF_HOJA.finditer(sin_lit)]


def celdas_de(sqref):
    s = set()
    for rng in str(sqref).split():
        c1, r1, c2, r2 = range_boundaries(rng)
        for r in range(r1, r2 + 1):
            for c in range(c1, c2 + 1):
                s.add((r, c))
    return s


def items_dv(f1):
    return f1.strip('"').split(',') if f1 and f1.startswith('"') else []


def claves_de_hoja(ws):
    s = set()
    for dv in ws.data_validations.dataValidation:
        s.update(items_dv(dv.formula1))
        if dv.type == 'custom':
            s.update(literales(dv.formula1))
    for row in ws.iter_rows():
        for c in row:
            if isinstance(c.value, str) and c.data_type == 'f':
                s.update(literales(c.value))
    for rng, reglas in ws.conditional_formatting._cf_rules.items():
        for rg in reglas:
            for f in rg.formula or []:
                s.update(literales(f))
    return s


def clasificar_celda(fichero, ws, cel, texto, hojas_libro, claves_hoja):
    f = corto(fichero)
    k = (f, ws.title, cel.coordinate)
    if k in mapas.POR_CELDA:
        return 'regenerar', 'rótulo fijado (mapas.POR_CELDA)'
    if k in mapas.VALORES_EN:
        return 'regenerar', 'valor fijado (mapas.VALORES_EN)'
    if ws.title == 'Instrucciones' and cel.coordinate == 'B2':
        return 'regenerar', 'D5 título interior (mapas.titulo_interior)'
    if RX_VERSION.match(texto):
        return 'regenerar', 'D19 línea de versión'
    if texto == mapas.MARCA_ES:
        return 'regenerar', 'D2 línea de marca «%s»' % mapas.MARCA_EN
    if mapas.RX_HORA_CELDA.match(texto):
        return 'regenerar', 'D9 hora de 12 h «%s» (mapas.hora_us)' % mapas.hora_us(texto)
    if texto in mapas.GLOSARIO:
        return 'mapa-glosario', 'mapas.GLOSARIO'
    if texto in hojas_libro:
        return 'mapa-hoja', 'celda que cita la pestaña'
    if texto in claves_hoja and texto in mapas.CLAVES:
        return 'mapa-clave', 'valor de DV/literal de la hoja'
    return 'traducir', None


PLANO = getattr(mapas, 'MOLDE', None) == 'plano'     # foodtruck/SPEC.md §0: una hoja por libro, sin columna Nº


def filas_tarea(ws):
    """Filas de tarea del molde PLANO (foodtruck/SPEC.md §0.1): por debajo de la cabecera (mapas.FILA_CABECERA), texto
    en A y en B (B sin fórmula) y A que no es el ancla de una banda fusionada (sección, título, línea de marca). El
    molde del kit base sigue con su regla (nº entero en A y texto en B, desde la fila 5)."""
    bandas = {(m.min_row, m.min_col) for m in ws.merged_cells.ranges if m.max_col > m.min_col}
    out = set()
    for r in range(mapas.FILA_CABECERA + 1, ws.max_row + 1):
        a, b = ws.cell(row=r, column=1), ws.cell(row=r, column=2)
        if (isinstance(a.value, str) and a.value.strip() and isinstance(b.value, str) and b.value.strip()
                and b.data_type != 'f' and (r, 1) not in bandas):
            out.add(r)
    return out


def dependencias_de(texto, entera=True):
    out = []
    for dep, (_etq, rx, _dest) in DEPENDENCIAS.items():
        if dep == 'hora24':
            if entera and mapas.RX_HORA_CELDA.match(texto):
                out.append(dep)
        elif dep == 'puestos':
            if entera and texto in mapas.PUESTOS_ES:
                out.append(dep)
        elif dep == 'hora_texto':
            if re.search(rx, texto) and not mapas.RX_HORA_CELDA.match(texto):
                out.append(dep)
        elif re.search(rx, texto):
            out.append(dep)
    return out


# --------------------------------------------------------------------------
# censo de un libro
# --------------------------------------------------------------------------
def censar_libro(path, fichero):
    wb = openpyxl.load_workbook(path)
    hojas = list(wb.sheetnames)
    ocurr = []
    info = OrderedDict()
    info['corto'] = corto(fichero)
    info['sha256'] = sha256(path)
    info['nombre_en'] = mapas.FICHEROS[fichero]
    info['titulo_en'] = mapas.TITULOS[mapas.FICHEROS[fichero]]
    info['hojas'] = hojas
    info['hojas_en'] = [mapas.HOJAS.get(h) for h in hojas]
    info['nombres_definidos'] = sorted(getattr(wb.defined_names, 'keys', lambda: [])())
    pr = wb.properties
    info['docprops'] = OrderedDict((k, getattr(pr, k)) for k in
                                   ('creator', 'title', 'subject', 'keywords', 'description', 'category',
                                    'lastModifiedBy'))
    info['hojas_detalle'] = OrderedDict()
    for h in hojas:
        ocurr.append((h, {'f': corto(fichero), 'h': h, 't': 'hoja', 'tr': 'mapa-hoja', 'det': 'nombre de pestaña'}))

    fx_hoja_total = 0
    claves_libro = set()
    for ws in wb.worksheets:
        claves_libro |= claves_de_hoja(ws)
    for ws in wb.worksheets:
        hd = OrderedDict()
        hd['nombre_en'] = mapas.HOJAS.get(ws.title)
        hd['dimension'] = ws.dimensions
        hd['max_fila'] = ws.max_row
        hd['max_col'] = ws.max_column
        n_txt = n_num = n_fx = fx_con_hoja = 0
        refs = defaultdict(int)
        lits = defaultdict(list)
        formatos = defaultdict(int)
        fechas, horas, numeros, celsius, moneda, fmt_euro = [], [], [], [], [], []
        textos, formulas = OrderedDict(), OrderedDict()
        verdes = verdes_desbl = desbl = 0
        patrones = OrderedDict()
        huella = hashlib.sha1()
        for row in ws.iter_rows():
            for cel in row:
                if cel.__class__.__name__ == 'MergedCell':
                    continue
                v = cel.value
                fmt = cel.number_format
                fid = getattr(cel._style, 'numFmtId', None)
                if is_date_format(fmt):
                    d = OrderedDict([('celda', cel.coordinate), ('numFmtId', fid), ('formato', fmt),
                                     ('vacia', v is None)])
                    (horas if RX_HORA_FMT.search(fmt) and not re.search(r'[dy]', fmt, re.I) else fechas).append(d)
                if fmt and fmt != 'General':
                    formatos[fmt] += 1
                if fmt and '€' in fmt:
                    fmt_euro.append(cel.coordinate)
                fill = cel.fill.fgColor.rgb if cel.fill is not None and cel.fill.fgColor is not None else None
                bloq = cel.protection is None or cel.protection.locked is not False
                if not bloq:
                    desbl += 1
                if isinstance(fill, str) and fill.upper().endswith(VERDE):
                    verdes += 1
                    if not bloq:
                        verdes_desbl += 1
                if v is None:
                    continue
                if isinstance(v, str) and cel.data_type == 'f':
                    n_fx += 1
                    formulas[cel.coordinate] = v
                    huella.update('{}={}\n'.format(cel.coordinate, v).encode('utf-8'))
                    for lit in literales(v):
                        lits[lit].append(cel.coordinate)
                    rh = refs_hoja(v)
                    if rh:
                        fx_con_hoja += 1
                    for nom in rh:
                        refs[nom] += 1
                    p = mapas.patron(v)
                    if p not in patrones:
                        patrones[p] = OrderedDict([('patron', p), ('n', 0), ('celdas', []),
                                                   ('cifras', bool(RX_COMPARA.search(RX_LITERAL.sub('""', v))))])
                    patrones[p]['n'] += 1
                    patrones[p]['celdas'].append(cel.coordinate)
                    continue
                if isinstance(v, str):
                    n_txt += 1
                    textos[cel.coordinate] = v
                    if '°C' in v:
                        celsius.append(cel.coordinate)
                    if re.search(DEPENDENCIAS['moneda'][1], v):
                        moneda.append(cel.coordinate)
                    tr, det = clasificar_celda(fichero, ws, cel, v, hojas, claves_libro)
                    d = {'f': corto(fichero), 'h': ws.title, 'c': cel.coordinate, 't': 'celda', 'tr': tr}
                    if det:
                        d['det'] = det
                    ad = mapas.ADAPTAR.get((corto(fichero), ws.title, cel.coordinate))
                    if ad:
                        d['adaptar'] = ad
                    ocurr.append((v, d))
                else:
                    n_num += 1
                    if ws.title != 'Instrucciones':
                        val = v.isoformat() if hasattr(v, 'isoformat') else v
                        numeros.append([cel.coordinate, val, fmt])
        fx_hoja_total += fx_con_hoja
        for p in patrones.values():
            cs = p['celdas']
            p['celdas'] = cs[0] + ('..' + cs[-1] if len(cs) > 1 else '')
        # tareas = filas con número en A y texto en B (las de los checklists; las vacías no cuentan)
        tareas = sum(1 for r in range(5, ws.max_row + 1)
                     if isinstance(ws.cell(row=r, column=1).value, int) and
                     isinstance(ws.cell(row=r, column=2).value, str) and ws.cell(row=r, column=2).data_type != 'f')
        if PLANO:                                                # molde plano: tareas sin número en A
            tareas = len(filas_tarea(ws))
        hd['tareas'] = tareas if ws.title != 'Instrucciones' else 0
        hd['celdas_texto'] = n_txt
        hd['celdas_numero'] = n_num
        hd['celdas_formula'] = n_fx
        hd['formulas_con_hoja'] = fx_con_hoja
        hd['refs_hoja'] = OrderedDict(sorted(refs.items()))
        hd['literales'] = OrderedDict((k, lits[k]) for k in sorted(lits))
        hd['literales_recuento'] = OrderedDict((k, len(lits[k])) for k in sorted(lits))
        hd['huella_formulas_sha1'] = huella.hexdigest() if n_fx else None
        hd['patrones_formula'] = list(patrones.values())
        hd['fechas'] = fechas
        hd['horas'] = horas
        hd['numeros'] = numeros
        hd['formatos'] = OrderedDict(sorted(formatos.items(), key=lambda kv: -kv[1]))
        hd['formatos_euro'] = fmt_euro
        hd['textos_celsius'] = celsius
        hd['textos_moneda'] = moneda
        hd['verdes'] = verdes
        hd['verdes_desbloqueadas'] = verdes_desbl
        hd['desbloqueadas'] = desbl
        hd['columnas_ocultas'] = sorted(k for k, d in ws.column_dimensions.items() if d.hidden)
        hd['anchos'] = OrderedDict((k, d.width) for k, d in sorted(ws.column_dimensions.items()) if d.width)
        hd['merges'] = sorted(str(m) for m in ws.merged_cells.ranges)
        hd['paneles'] = ws.freeze_panes
        p = ws.protection
        hd['proteccion'] = OrderedDict([('hoja', bool(p.sheet)), ('password', bool(p.password)),
                                        ('formatColumns', p.formatColumns), ('formatRows', p.formatRows),
                                        ('insertRows', p.insertRows), ('deleteRows', p.deleteRows),
                                        ('sort', p.sort), ('autoFilter', p.autoFilter)])
        hd['autofiltro'] = ws.auto_filter.ref
        hd['area_impresion'] = ws.print_area
        hd['titulos_impresion'] = OrderedDict([('filas', ws.print_title_rows), ('columnas', ws.print_title_cols)])
        ps = ws.page_setup
        hd['papel'] = OrderedDict([('paperSize', ps.paperSize), ('orientation', ps.orientation),
                                   ('fitToWidth', ps.fitToWidth), ('fitToHeight', ps.fitToHeight),
                                   ('fitToPage', bool(ws.sheet_properties.pageSetUpPr and
                                                      ws.sheet_properties.pageSetUpPr.fitToPage))])
        hf = OrderedDict()
        for nombre in ('oddHeader', 'oddFooter', 'evenHeader', 'evenFooter', 'firstHeader', 'firstFooter'):
            obj = getattr(ws, nombre)
            for parte in ('left', 'center', 'right'):
                txt = getattr(obj, parte).text
                if txt:
                    hf['{}.{}'.format(nombre, parte)] = txt
                    ocurr.append((txt, {'f': corto(fichero), 'h': ws.title, 't': 'cabecera-pie',
                                        'c': '{}.{}'.format(nombre, parte), 'tr': 'regenerar',
                                        'det': 'D19 pie «%s»' % mapas.FOOTER_EN}))
        hd['cabeceras_pies'] = hf
        dvs = []
        for dv in ws.data_validations.dataValidation:
            dvs.append(OrderedDict([
                ('sqref', str(dv.sqref)), ('type', dv.type), ('operator', dv.operator),
                ('formula1', dv.formula1), ('formula2', dv.formula2), ('allow_blank', dv.allow_blank),
                ('showErrorMessage', dv.showErrorMessage), ('showInputMessage', dv.showInputMessage),
                ('errorStyle', dv.errorStyle),
                ('errorTitle', dv.errorTitle), ('error', dv.error),
                ('promptTitle', dv.promptTitle), ('prompt', dv.prompt),
                ('n_celdas', len(celdas_de(dv.sqref))),
                ('cita_hojas', sorted(set(refs_hoja(dv.formula1 or '')))),
                ('literales', literales(dv.formula1) if dv.type == 'custom' else []),
            ]))
            base = {'f': corto(fichero), 'h': ws.title, 'c': str(dv.sqref)}
            for campo, t in (('errorTitle', 'dv-error-titulo'), ('error', 'dv-error'),
                             ('promptTitle', 'dv-prompt-titulo'), ('prompt', 'dv-prompt')):
                txt = getattr(dv, campo)
                if txt:
                    ocurr.append((txt, dict(base, t=t, tr='traducir')))
            for it in items_dv(dv.formula1):
                if it == '':
                    continue
                ocurr.append((it, dict(base, t='dv-item', tr='mapa-clave', det='mapas.CLAVES')))
        hd['dv'] = dvs
        cfs = []
        lits_cf = defaultdict(int)
        for rng, reglas in ws.conditional_formatting._cf_rules.items():
            for rg in reglas:
                fs = list(rg.formula or [])
                cfs.append(OrderedDict([
                    ('sqref', str(rng.sqref)), ('type', rg.type), ('operator', rg.operator),
                    ('text', getattr(rg, 'text', None)), ('formula', fs), ('priority', rg.priority),
                    ('stopIfTrue', rg.stopIfTrue), ('cita_hojas', sorted({h for f in fs for h in refs_hoja(f)})),
                    ('tokens', [t for f in fs for t in literales(f)]),
                    ('fill', getattr(getattr(getattr(rg.dxf, 'fill', None), 'fgColor', None), 'rgb', None)
                     if rg.dxf is not None else None),
                ]))
                for f in fs:
                    for lit in literales(f):
                        lits_cf[lit] += 1
        hd['cf'] = cfs
        hd['literales_cf'] = OrderedDict(sorted(lits_cf.items()))
        hd['textos'] = textos
        hd['formulas'] = formulas
        info['hojas_detalle'][ws.title] = hd

    with zipfile.ZipFile(path) as z:
        nombres = z.namelist()
        info['graficos'] = sorted(n for n in nombres if re.fullmatch(r'xl/charts/chart\d+\.xml', n))
        info['imagenes'] = sorted(n for n in nombres if n.startswith('xl/media/'))
        info['comentarios'] = sorted(n for n in nombres if re.fullmatch(r'xl/comments\d*\.xml', n))
    for campo, v in info['docprops'].items():
        if not isinstance(v, str) or not v or campo == 'creator':
            continue
        ocurr.append((v, {'f': corto(fichero), 't': 'docprops', 'c': campo, 'tr': 'regenerar',
                          'det': 'D19 docProps'}))

    hdv = list(info['hojas_detalle'].values())
    lit_tot, litcf_tot = defaultdict(int), defaultdict(int)
    for x in hdv:
        for k, v in x['literales_recuento'].items():
            lit_tot[k] += v
        for k, v in x['literales_cf'].items():
            litcf_tot[k] += v
    info['totales'] = OrderedDict([
        ('hojas', len(hojas)),
        ('celdas_texto', sum(x['celdas_texto'] for x in hdv)),
        ('celdas_formula', sum(x['celdas_formula'] for x in hdv)),
        ('patrones_formula', sum(len(x['patrones_formula']) for x in hdv)),
        ('formulas_con_hoja', fx_hoja_total),
        ('dv', sum(len(x['dv']) for x in hdv)),
        ('dv_lista_literal', sum(1 for x in hdv for d in x['dv'] if (d['formula1'] or '').startswith('"'))),
        ('dv_celdas', sum(d['n_celdas'] for x in hdv for d in x['dv'])),
        ('cf', sum(len(x['cf']) for x in hdv)),
        ('graficos', len(info['graficos'])),
        ('imagenes', len(info['imagenes'])),
        ('comentarios', len(info['comentarios'])),
        ('merges', sum(len(x['merges']) for x in hdv)),
        ('hojas_protegidas', sum(1 for x in hdv if x['proteccion']['hoja'])),
        ('hojas_con_password', sum(1 for x in hdv if x['proteccion']['password'])),
        ('areas_impresion', sum(1 for x in hdv if x['area_impresion'])),
        ('titulos_impresion', sum(1 for x in hdv if x['titulos_impresion']['filas'])),
        ('paneles', sum(1 for x in hdv if x['paneles'])),
        ('tareas', sum(x['tareas'] for x in hdv)),
        ('fechas', sum(len(x['fechas']) for x in hdv)),
        ('horas', sum(len(x['horas']) for x in hdv)),
        ('numeros_ejemplo', sum(len(x['numeros']) for x in hdv)),
        ('formatos_euro', sum(len(x['formatos_euro']) for x in hdv)),
        ('textos_celsius', sum(len(x['textos_celsius']) for x in hdv)),
        ('textos_moneda', sum(len(x['textos_moneda']) for x in hdv)),
        ('horas_24h', sum(1 for x in hdv for t in x['textos'].values() if mapas.RX_HORA_CELDA.match(t))),
        ('verdes', sum(x['verdes'] for x in hdv)),
        ('verdes_desbloqueadas', sum(x['verdes_desbloqueadas'] for x in hdv)),
        ('desbloqueadas', sum(x['desbloqueadas'] for x in hdv)),
        ('columnas_ocultas', sum(len(x['columnas_ocultas']) for x in hdv)),
        ('papel', sorted({x['papel']['paperSize'] for x in hdv if x['papel']['paperSize'] is not None})),
        ('literales', OrderedDict(sorted(lit_tot.items()))),
        ('literales_cf', OrderedDict(sorted(litcf_tot.items()))),
    ])
    return info, ocurr


# --------------------------------------------------------------------------
# cadenas, grupos y pistas
# --------------------------------------------------------------------------
PRIORIDAD = {'traducir': 0, 'localizar': 1, 'mapa-hoja': 2, 'mapa-clave': 3, 'mapa-glosario': 4, 'regenerar': 5}
SUMABLES = ('hojas', 'celdas_texto', 'celdas_formula', 'patrones_formula', 'formulas_con_hoja', 'dv',
            'dv_lista_literal', 'dv_celdas', 'cf', 'graficos', 'imagenes', 'comentarios',
            'merges', 'hojas_protegidas', 'hojas_con_password', 'areas_impresion', 'titulos_impresion', 'paneles',
            'tareas', 'fechas', 'horas', 'numeros_ejemplo', 'formatos_euro', 'textos_celsius', 'textos_moneda',
            'horas_24h', 'verdes', 'verdes_desbloqueadas', 'desbloqueadas', 'columnas_ocultas')


def pistas(texto):
    return [nota for rx, nota in PISTAS if re.search(rx, texto)]


def citas(texto, citables):
    out = []
    for m in RX_CITA.finditer(texto):
        span = next(g for g in m.groups() if g is not None)
        if span in citables and span not in [c['es'] for c in out]:
            out.append(OrderedDict([('es', span), ('en', citables[span])]))
    return out


def citas_fichero(texto):
    out = []
    for m in RX_FICHERO.finditer(texto):
        es = m.group(0)
        if es in mapas.FICHEROS and es not in [c['es'] for c in out]:
            out.append(OrderedDict([('es', es), ('en', mapas.FICHEROS[es])]))
    return out


def palabras(s):
    return len(re.findall(r'\w+', s))


def dos_bloques(orden, peso):
    """Parte la lista de libros en 2 bloques contiguos minimizando el bloque más pesado (G1 lleva las comunes)."""
    mejor = None
    for i in range(1, len(orden)):
        bl = [orden[:i], orden[i:]]
        m = max(sum(peso[f] for f in b) + (peso['G1-comun'] if k == 0 else 0) for k, b in enumerate(bl))
        if mejor is None or m < mejor[0]:
            mejor = (m, bl)
    return mejor[1]


def main():
    ap = argparse.ArgumentParser(description=__doc__.split('\n')[1])
    kit.argumento(ap)
    ap.add_argument('--origen', default=ORIGEN)
    ap.add_argument('--salida', default=kit.DATOS)
    ap.add_argument('--md', action='store_true', help='imprime las tablas §0 y §1 de F1-inventario-es.md')
    args = ap.parse_args()

    libros = list(mapas.LIBROS_ES)
    faltan = [f for f in libros if not os.path.exists(os.path.join(args.origen, f))]
    if faltan:
        sys.exit('faltan ficheros en {}: {}'.format(args.origen, faltan))
    sobran = sorted(f for f in os.listdir(args.origen) if f.endswith('.xlsx') and f not in libros)
    if sobran:
        sys.exit('xlsx no previstos en {}: {}'.format(args.origen, sobran))
    try:
        commit = subprocess.check_output(['git', '-C', REPO, 'rev-parse', '--short', 'HEAD'],
                                         stderr=subprocess.DEVNULL).decode().strip()
    except Exception:                                             # noqa: BLE001
        commit = None

    censo = OrderedDict()
    censo['meta'] = OrderedDict([
        ('generado', datetime.datetime.utcnow().replace(microsecond=0).isoformat() + 'Z'),
        ('script', 'scripts/productos-digitales/checklists-kit/extraer_textos.py'),
        ('origen', os.path.relpath(args.origen, REPO)),
        ('commit', commit),
        ('nota', 'Censo del Kit de Tareas Recurrentes v2.0 PUBLICADO (11 xlsx). Los gates de la F2-EN comparan '
                 'contra este fichero, nunca contra recuentos fijos. patrones_formula = fórmulas con las filas '
                 'normalizadas a {r} (mapas.patron); «cifras» = el patrón compara con un número. tareas = filas '
                 'con número en A y texto en B. dependencias = lo que depende de España (F1-inventario §1).'),
    ])
    censo['libros'] = OrderedDict()
    todas = []
    for f in libros:
        info, oc = censar_libro(os.path.join(args.origen, f), f)
        censo['libros'][f] = info
        todas.extend(oc)
    T = OrderedDict((k, sum(censo['libros'][f]['totales'][k] for f in libros)) for k in SUMABLES)
    lit, litcf = defaultdict(int), defaultdict(int)
    for f in libros:
        for k, v in censo['libros'][f]['totales']['literales'].items():
            lit[k] += v
        for k, v in censo['libros'][f]['totales']['literales_cf'].items():
            litcf[k] += v
    T['literales'] = OrderedDict(sorted(lit.items(), key=lambda kv: -kv[1]))
    T['literales_cf'] = OrderedDict(sorted(litcf.items(), key=lambda kv: -kv[1]))
    T['papel'] = sorted({p for f in libros for p in censo['libros'][f]['totales']['papel']})
    censo['totales'] = T

    # ---- dependencias de España (por aparición en celda o mensaje de DV; pies y docProps aparte)
    dep = OrderedDict((k, OrderedDict([('etiqueta', v[0]), ('destino', v[2]), ('apariciones', 0),
                                       ('cadenas', 0), ('donde', [])])) for k, v in DEPENDENCIAS.items())
    vistas = defaultdict(set)
    for texto, oc in todas:
        if oc['t'] not in ('celda', 'dv-error', 'dv-error-titulo', 'dv-prompt', 'dv-prompt-titulo'):
            continue
        for d in dependencias_de(texto, entera=(oc['t'] == 'celda')):
            dep[d]['apariciones'] += 1
            vistas[d].add(texto)
            if len(dep[d]['donde']) < 400:
                dep[d]['donde'].append('%s:%s!%s' % (oc['f'], oc['h'], oc['c']))
    for d in dep:
        dep[d]['cadenas'] = len(vistas[d])
    dep['fmt_euro'] = OrderedDict([('etiqueta', 'Formatos de número con €'), ('destino', 'D13: `#,##0.00`'),
                                   ('apariciones', T['formatos_euro']), ('cadenas', 0), ('donde', [])])
    dep['papel'] = OrderedDict([('etiqueta', 'Hojas en A4 (paperSize 9)'), ('destino', 'D17: US Letter (paperSize 1)'),
                                ('apariciones', sum(1 for f in libros for h in censo['libros'][f]['hojas_detalle']
                                                    .values() if h['papel']['paperSize'] == 9)),
                                ('cadenas', 0), ('donde', [])])
    censo['dependencias'] = dep

    # ---- citables: pestañas, glosario y rótulos fijados (ES → EN)
    citables = OrderedDict()
    for f in libros:
        for h in censo['libros'][f]['hojas']:
            citables[h] = mapas.HOJAS.get(h)
    for es, en in mapas.GLOSARIO.items():
        citables.setdefault(es, en)
    for (c, h, celda), en in mapas.POR_CELDA.items():
        info = next(i for i in censo['libros'].values() if i['corto'] == c)
        es = info['hojas_detalle'][h]['textos'].get(celda)
        if es:
            citables.setdefault(es, en)
            citables.setdefault(re.sub(r'^\W+\s*', '', es), re.sub(r'^\W+\s*', '', en))

    # ---- cadenas únicas
    excl = OrderedDict()
    cad = OrderedDict()
    for texto, oc in todas:
        if (not tiene_letra(texto) and not re.search(r'€|°C', texto)) or texto in INVARIABLES:
            excl[texto] = excl.get(texto, 0) + 1
            continue
        cad.setdefault(texto, []).append(oc)
    orden_f = [corto(f) for f in libros]
    cadenas = []
    for i, (texto, ocs) in enumerate(cad.items(), 1):
        fich = []
        for o in ocs:
            if o['f'] not in fich:
                fich.append(o['f'])
        trs = sorted({o['tr'] for o in ocs}, key=lambda t: PRIORIDAD[t])
        e = OrderedDict()
        e['id'] = 'c{:04d}'.format(i)
        e['es'] = texto
        e['grupo'] = None
        e['tratamiento'] = trs[0]
        if len(trs) > 1:
            e['tratamientos'] = trs
        if 'mapa-hoja' in trs and texto in mapas.HOJAS:
            e['en_hoja'] = mapas.HOJAS[texto]
        if 'mapa-glosario' in trs:
            e['en_glosario'] = mapas.GLOSARIO[texto]
        if 'mapa-clave' in trs and texto in mapas.CLAVES:
            e['en_clave'] = mapas.CLAVES[texto]
        if trs[0] == 'regenerar':
            pc = [mapas.POR_CELDA[(o['f'], o['h'], o['c'])] for o in ocs if (o.get('f'), o.get('h'), o.get('c'))
                  in mapas.POR_CELDA]
            if pc:
                e['en_por_celda'] = pc[0]
        e['n_ficheros'] = len(fich)
        e['n_apariciones'] = len(ocs)
        e['palabras'] = palabras(texto)
        e['con_cifra'] = bool(re.search(r'\d', texto))
        ch = citas(texto, citables)
        if ch:
            e['cita_rotulos'] = ch
        cf = citas_fichero(texto)
        if cf:
            e['cita_ficheros'] = cf
        ads = [o['adaptar'] for o in ocs if o.get('adaptar')]
        if ads:
            e['adaptar'] = ads[0] if len(set(ads)) == 1 else ' | '.join(sorted(set(ads)))
        if trs[0] in ('traducir', 'localizar'):
            pi = pistas(texto)
            if pi:
                e['pistas'] = pi
        e['sha1'] = hashlib.sha1(texto.encode('utf-8')).hexdigest()[:12]
        e['donde'] = [{k: v for k, v in o.items() if k != 'adaptar'} for o in ocs]
        e['_ficheros'] = fich
        cadenas.append(e)

    # ---- grupos: G1-G2 (subagentes) por bloques contiguos de libros; comunes → G1; GM no se traduce
    peso = defaultdict(int)
    for e in cadenas:
        if e['tratamiento'] in ('traducir', 'localizar'):
            comun = e['n_ficheros'] >= 3
            e['_comun'] = comun
            peso['G1-comun' if comun else e['_ficheros'][0]] += e['palabras']
    bloques = dos_bloques(orden_f, peso)
    gid_de = {}
    for k, bl in enumerate(bloques, 1):
        for f in bl:
            gid_de[f] = 'G%d' % k
    for e in cadenas:
        if e['tratamiento'] in ('traducir', 'localizar'):
            e['grupo'] = 'G1' if e['_comun'] else gid_de[e['_ficheros'][0]]
            if e['_comun']:
                e['comun'] = True
        else:
            e['grupo'] = 'GM'
    grupos = []
    for k, bl in enumerate(bloques, 1):
        gid = 'G%d' % k
        es = [e for e in cadenas if e['grupo'] == gid]
        grupos.append(OrderedDict([
            ('id', gid),
            ('descripcion', ('Cadenas comunes a ≥ 3 libros (se traducen PRIMERO: fijan el glosario) + ' if k == 1 else '')
             + 'libros ' + ', '.join(bl)),
            ('libros', bl),
            ('n_cadenas', len(es)), ('n_adaptar', sum(1 for e in es if e.get('adaptar'))),
            ('n_palabras', sum(e['palabras'] for e in es)),
        ]))
    gm = [e for e in cadenas if e['grupo'] == 'GM']
    por_tr = defaultdict(int)
    for e in gm:
        por_tr[e['tratamiento']] += 1
    grupos.append(OrderedDict([
        ('id', 'GM'),
        ('descripcion', 'NO se traduce: pestañas, claves y glosario (mapas.py), horas, rótulos fijados, versión, '
                        'marca, pie, docProps. ' + ', '.join('{} {}'.format(v, k) for k, v in sorted(por_tr.items()))),
        ('n_cadenas', len(gm)), ('n_palabras', sum(e['palabras'] for e in gm)),
    ]))
    orden_g = {g['id']: k for k, g in enumerate(grupos)}
    cadenas.sort(key=lambda e: (orden_g[e['grupo']], int(e['id'][1:])))
    for e in cadenas:
        e.pop('_ficheros')
        e.pop('_comun', None)
    assert all(e['grupo'] for e in cadenas)

    textos = OrderedDict()
    textos['meta'] = OrderedDict([
        ('generado', censo['meta']['generado']),
        ('script', censo['meta']['script']),
        ('origen', censo['meta']['origen']),
        ('commit', commit),
        ('libros', OrderedDict((corto(f), OrderedDict([('es', f), ('en', mapas.FICHEROS[f]),
                                                      ('sha256', censo['libros'][f]['sha256'])]))
                               for f in libros)),
        ('n_cadenas', len(cadenas)),
        ('n_traducir', sum(1 for e in cadenas if e['tratamiento'] == 'traducir')),
        ('n_adaptar', sum(1 for e in cadenas if e.get('adaptar') and e['grupo'] != 'GM')),
        ('n_palabras_subagentes', sum(e['palabras'] for e in cadenas if e['grupo'] != 'GM')),
        ('n_apariciones', sum(e['n_apariciones'] for e in cadenas)),
        ('grupos', grupos),
        ('excluidas', OrderedDict([('cadenas', excl),
                                   ('regla', 'Fuera: cadenas sin letras (salvo las que llevan € o °C) y las invariables %s. Los literales de '
                                             'fórmula están en censo_es.json (mapas.CLAVES).' % sorted(INVARIABLES))])),
        ('campos', OrderedDict([
            ('donde', 'f = libro, h = hoja, c = celda / sqref / campo; t = celda | hoja | dv-* | cabecera-pie | '
                      'docprops; tr = tratamiento de esa aparición; det = detalle'),
            ('en_hoja / en_glosario / en_clave / en_por_celda', 'valor EN obligatorio que dicta mapas.py'),
            ('cita_rotulos', 'pestañas, rótulos y valores citados entre comillas, con su EN obligatorio'),
            ('cita_ficheros', 'ficheros del kit citados en el texto, con su nombre EN'),
            ('adaptar', 'instrucción de adaptación al mercado US para esa celda (mapas.ADAPTAR): manda sobre la '
                        'traducción literal'),
            ('pistas', 'decisiones de la SPEC que aplican a la cadena'),
        ])),
    ])
    textos['cadenas'] = cadenas

    os.makedirs(args.salida, exist_ok=True)
    for nombre, obj in (('textos_es.json', textos), ('censo_es.json', censo)):
        with open(os.path.join(args.salida, nombre), 'w', encoding='utf-8') as fh:
            json.dump(obj, fh, ensure_ascii=False, indent=1)
            fh.write('\n')

    print('textos_es.json: {} cadenas ({} traducir, {} con ADAPTAR, {} palabras para subagentes) · {} apariciones'
          .format(len(cadenas), textos['meta']['n_traducir'], textos['meta']['n_adaptar'],
                  textos['meta']['n_palabras_subagentes'], textos['meta']['n_apariciones']))
    for g in grupos:
        print('  {:<3} {:>4} cadenas {:>6} palabras  {}'.format(g['id'], g['n_cadenas'], g['n_palabras'],
                                                              g['descripcion'][:110]))
    print('censo_es.json: ' + json.dumps({k: v for k, v in T.items() if not isinstance(v, dict)},
                                         ensure_ascii=False))
    if args.md:
        print(tabla_md(censo))
        print()
        print(tabla_dependencias_md(censo))


COLS_MD = [('hojas', 'Hojas'), ('celdas_texto', 'Texto'), ('tareas', 'Tareas'), ('celdas_formula', 'Fórmulas'),
           ('formulas_con_hoja', 'Fx→hoja'), ('dv', 'DV'), ('cf', 'CF'), ('merges', 'Merges'),
           ('hojas_protegidas', 'Protegidas'), ('horas_24h', 'Horas 24 h'), ('textos_celsius', '°C'),
           ('formatos_euro', 'Fmt €'), ('textos_moneda', 'Textos €'), ('verdes', 'Verdes'),
           ('desbloqueadas', 'Desbloq.')]


def tabla_md(censo):
    out = ['| Libro | ' + ' | '.join(t for _, t in COLS_MD) + ' |', '|---|' + '---|' * len(COLS_MD)]
    for f, info in censo['libros'].items():
        out.append('| %s | ' % f[:-5] + ' | '.join(str(info['totales'][k]) for k, _ in COLS_MD) + ' |')
    out.append('| **Total** | ' + ' | '.join('**%s**' % censo['totales'][k] for k, _ in COLS_MD) + ' |')
    return '\n'.join(out)


def tabla_dependencias_md(censo):
    out = ['| Dependencia de España | Apariciones | Cadenas | Destino |', '|---|---|---|---|']
    for k, d in censo['dependencias'].items():
        out.append('| %s | %d | %s | %s |' % (d['etiqueta'], d['apariciones'], d['cadenas'] or '—', d['destino']))
    return '\n'.join(out)


if __name__ == '__main__':
    main()
