#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
gates_en.py — Gates F2 del Restaurant Checklists Kit Pro (EN), SPEC §8 (G1-G8), sobre la salida de aplicar_en.py.
Copia adaptada de `staff-kit/gates_en.py` (sesión Claude Code, 17-oct-2026). Solo LEE: los 11 xlsx EN, los 11 xlsx
ES publicados (dl/kit-tareas, solo lectura), censo_es.json, textos_es.json, textos_en/ y mapas.py. Ningún recuento
está escrito a mano: todo sale de censo_es.json, de los xlsx ES o de mapas.py; las únicas cifras propias son las que
la SPEC fija (§5 caso de caja y §8 G7, que además se cruzan con mapas.py).

    python3 gates_en.py                               # contra astro-site/public/dl/restaurant-checklist-templates/
    python3 gates_en.py --kit bar ...                 # un vertical (kit.py), contra su censo y dl/<su slug>/
    python3 gates_en.py --dir /tmp/ck-dry             # contra un dry-run
    python3 gates_en.py --solo G1,G2 [--mes October] [--json informe.json]
    python3 gates_en.py --parcial [--textos DIR]      # dry-run parcial: la marca «PENDING-EN cNNNN» se tolera
    python3 gates_en.py --autotest [--dir …]          # inyecta defectos en una COPIA y exige que se detecten

Gates (SPEC §8):
  G1 paridad: 47 hojas por §2.2; fórmula a fórmula = ES tras HOJAS + CLAVES salvo las 15 FORMULAS_POR_CELDA;
     fórmulas con hoja; DV (sqref, nº de celdas, flags, tipo, ítems por CLAVES); CF (sqref, tipo, prioridad,
     relleno, fórmula); merges; áreas y títulos de impresión (también en el definedName crudo de workbook.xml:
     'Valentine''s Day'!); paneles; anchos; protección (36 hojas, sin contraseña, mismos flags); verdes y
     desbloqueadas celda a celda; las 491 tareas en las mismas filas; 0 gráficos e imágenes;
  G2 restos: cero español, «€», caracteres fuera de la lista blanca, coma decimal, comillas curvas, IDs internos,
     nombres de fichero ES y marca de relleno en valores, literales, DV, CF, formatos, pies y docProps; «°C» solo
     entre paréntesis tras su °F; ninguna celda «HH:MM» de 24 h; «$»/«¢» solo en la tabla de billetes (09
     A27:A42) y en las celdas cuyo ADAPTAR los pide; los formatos con € del censo reescritos; DV ≤ 32 / ≤ 255;
  G3 formato: paperSize 1 en las 47; pie, marca, versión y docProps D19; Instructions!B2 = título interior D5;
     ajuste de texto y alto de fila donde el EN ya no cabe (aplicar_en.objetivo_ajuste), sin desbordes;
  G4 mapas: cada aparición traducir / mapa-glosario / mapa-hoja / mapa-clave / regenerar = su EN (textos_en o
     mapas.py); horas = mapas.hora_us (las del censo); pestañas, rótulos y ficheros citados (cita_rotulos,
     cita_ficheros) con su EN exacto; firma de autor literal (SPEC §7.2);
  G5 caja: C27:C42 = mapas.DENOMINACIONES (fórmulas y rótulos) y el caso de SPEC §5 calculado en una copia
     (1 de cada → 204.91; cierre 200 / 1,250 / 3,000 / 150 / Z 4,200 → 4,200 y OVER / SHORT 0, en el cierre y
     en el registro mensual);
  G6 cálculo: <v> en toda <f>, cero errores; caché EN = caché ES (contadores incluidos);
  G7 reglas: cada °F de los textos ∈ SPEC §8 (= las cifras °F de mapas.py); cada cita normativa EN ∈ fuentes de
     mapas.REGLAS_US (SPEC §3); toda celda ES con °C lleva °F en el EN;
  G8 censo-entregables.py --fail (Letter) y gate-no-latinos.py en 0.
Sale con código 1 si algún gate falla.
"""
import argparse
import html
import json
import logging
import os
import re
import shutil
import subprocess
import sys
import tempfile
import unicodedata
import zipfile
from collections import OrderedDict, Counter

AQUI = os.path.dirname(os.path.abspath(__file__))
SCRIPTS = os.path.dirname(AQUI)
REPO = os.path.dirname(os.path.dirname(SCRIPTS))
sys.path.insert(0, AQUI)

import openpyxl                                          # noqa: E402

import kit                                               # noqa: E402  (--kit base|bar: kit.py)
from kit import mapas                                    # noqa: E402
import aplicar_en as A                                   # noqa: E402

X = A.extraer_textos
logging.disable(logging.CRITICAL)

ORIGEN = A.ORIGEN
CENSO_SCRIPT = os.path.join(SCRIPTS, 'censo-entregables.py')
NO_LATINOS = os.path.join(SCRIPTS, 'gate-no-latinos.py')
TOL = 0.006
ERRORES = ('#REF!', '#VALUE!', '#NAME?', '#DIV/0!', '#N/A', '#NUM!', '#NULL!')
CLAVES = mapas.CLAVES
CLAVES_EN = set(CLAVES.values())
FIRMA_EN = ('Designed by John Guerrero — chef and restaurant consultant since 2010, in professional kitchens since '
            'age 17 · johnguerrero.es')                                       # SPEC §7.2, literal
GRADOS_F_SPEC = set(getattr(mapas, 'GRADOS_F_SPEC', {-31, -4, 0, 10, 32, 41, 70, 100, 135, 165}))  # SPEC §8 G7
CAJA_SPEC = {'total_1_de_cada': 204.91, 'fondo': 200, 'contado': 1250, 'tarjetas': 3000, 'otros': 150, 'z': 4200}
# vertical con el 09 como libro propio (catering: cobros por evento): la tabla de billetes vive en mapas.HOJA_CAJA,
# G5 calcula mapas.CASOS_G5 y G1 admite las fórmulas con hoja que el EN añade (mapas.FORMULAS_CON_HOJA_EXTRA)
HOJA_CAJA = getattr(mapas, 'HOJA_CAJA', ('09', 'Cierre de Caja'))
CASOS_G5 = getattr(mapas, 'CASOS_G5', None)
CON_HOJA_EXTRA = getattr(mapas, 'FORMULAS_CON_HOJA_EXTRA', {})
SIN_MARCA = frozenset(getattr(mapas, 'LIBROS_SIN_MARCA', ()))   # libros cuyo ES no trae la línea de marca D19

SIGNOS_ES = set('¿¡«»€ºª')
RX_ID_INTERNO = re.compile(r'\[(?:fuente|estimado|derivado|criterio|local)[^\]]*\]|\bSPEC\b|\((?:D|I)\d{1,2}\)|'
                           r'\bD\d{1,2}\b(?=[:,)])|\bTEC-\d+\b|\bkit-tareas\b')
RX_CITA_ANIDADA = re.compile(r'\(\(|\)\)\)')
LIMITE_DV = {'promptTitle': 32, 'errorTitle': 32, 'prompt': 255, 'error': 255}     # Excel
RX_TILDE = re.compile(r'[áéíóúüñÁÉÍÓÚÜÑ]')
# «café» como palabra inglesa va sin tilde («sidewalk cafe permit», grafía de los permisos US): no entra.
LISTA_BLANCA_TILDE = {'entrée', 'entrées', 'sauté', 'sautéed', 'sautéing', 'purée', 'purées',
                      'jalapeño', 'jalapeños', 'piña', 'crème', 'brûlée', 'résumé', 'naïve', 'flambé', 'frappé',
                      'rosé', 'canapé', 'canapés', 'soufflé', 'consommé', 'crêpe', 'crêpes', 'déjà', 'à'}
FRASES_OK = ('Cinco de Mayo', 'al dente', 'pico de gallo', 'chili con carne', 'à la carte', 'a la carte',
             'Mise en Place', 'mise en place', 'sous vide', 'Sous Vide')
FUNCIONALES_ES = {
    'de', 'del', 'las', 'el', 'los', 'y', 'que', 'con', 'para', 'por', 'una', 'unos', 'unas', 'es', 'al', 'se',
    'su', 'sus', 'como', 'pero', 'muy', 'cuando', 'donde', 'porque', 'este', 'esta', 'esto', 'estos', 'desde',
    'hasta', 'sobre', 'entre', 'tu', 'tus', 'te', 'mis', 'hoja', 'pestaña', 'celda', 'celdas', 'sin', 'ejemplo',
    'cada', 'debe', 'hay', 'puede', 'si', 'lo', 'le', 'les', 'todo', 'toda', 'todos', 'nunca', 'siempre', 'ya',
}
OFICIO_ES = {
    'apertura', 'cierre', 'cocina', 'sala', 'barra', 'tarea', 'tareas', 'limpieza', 'camarero', 'camareros',
    'camarera', 'cocinero', 'cocineros', 'encargado', 'encargada', 'responsable', 'firma', 'notas', 'semanal',
    'mensual', 'anual', 'trimestral', 'caja', 'fondo', 'efectivo', 'recuento', 'descuadre', 'tarjeta', 'tarjetas',
    'ventas', 'mesas', 'carta', 'platos', 'nevera', 'camara', 'camaras', 'congelador', 'temperatura', 'aceite',
    'suelo', 'suelos', 'banos', 'terraza', 'puerta', 'persiana', 'persianas', 'luces', 'alarma', 'basura', 'cubos',
    'cubiertos', 'servilletas', 'copas', 'vasos', 'hielo', 'reservas', 'cliente', 'clientes', 'turno', 'turnos',
    'manana', 'tarde', 'noche', 'hoy', 'dia', 'dias', 'semana', 'semanas', 'fecha', 'hora', 'horas', 'partida',
    'partidas', 'plancha', 'freidora', 'fregadero', 'lavavajillas', 'campana', 'cuchillos', 'tablas', 'jefe',
    'auxiliar', 'gerente', 'equipo', 'cumpleanos', 'evento', 'eventos', 'festivo', 'festivos', 'navidad',
    'temporada', 'plantilla', 'plantillas', 'personalizable', 'negocio', 'recepcion', 'zona', 'zonas',
    'completada', 'completadas', 'verificado', 'instrucciones', 'restaurante', 'hosteleria', 'proveedor',
    'proveedores', 'pedido', 'pedidos', 'mermas', 'comanda', 'comandas', 'incidencias', 'resumen', 'registro',
    'concepto', 'importe', 'billetes', 'monedas', 'cantidad', 'denominacion', 'deposito', 'anotar', 'revisar',
    'comprobar', 'limpiar', 'preparar', 'encender', 'apagar', 'cerrar', 'abrir', 'verificar', 'marca', 'marcar',
}
RX_URL = re.compile(r'https?://\S+|\b[\w-]+(?:\.[\w-]+)*\.(?:es|com|pro|gov|org|uk|co|example|net)\b[/\w.#-]*', re.I)
RX_PALABRA = re.compile(r"[A-Za-zÀ-ÖØ-öø-ÿ]+")
RX_STR = A.RX_STR
RX_SIGLAS_ES = re.compile(r'\b(?:TPV|APPCC|IVA|CIF|NIF|DDD|BIE)\b')
RX_HORA_24 = re.compile(r'\b(?:0\d|1[3-9]|2[0-3]):[0-5]\d\b')
RX_COMA_DECIMAL = re.compile(r'(?<![\d,])\d+,\d{1,2}(?![\d])')
RX_FICHERO_ES = re.compile('|'.join(re.escape(f) for f in mapas.FICHEROS))
COMILLAS_CURVAS = set(A.COMILLAS)


def _simbolos_ok():
    """Lista blanca de caracteres no ASCII que no son letras: los del ES publicado (▸ ✓ — ☐ 💰 📊 …) + la
    tipografía del kit EN (– − ° · ¢). Los signos españoles (¿ ¡ « » € º ª) nunca entran."""
    ok = set('▸→☐✓✗⚠·—–…°§≤≥−±×½✅❌⭐¢') | {'️'}
    te = json.load(open(os.path.join(kit.DATOS, 'textos_es.json'), encoding='utf-8'))
    textos = [c['es'] for c in te['cadenas']] + list(te['meta']['excluidas']['cadenas'])
    censo = json.load(open(os.path.join(kit.DATOS, 'censo_es.json'), encoding='utf-8'))
    for info in censo['libros'].values():
        for hd in info['hojas_detalle'].values():
            textos += list(hd['literales']) + list(hd['literales_cf'])
    for t in textos:
        for ch in t:
            if ord(ch) > 127 and not unicodedata.category(ch).startswith('L') and ch not in SIGNOS_ES:
                ok.add(ch)
    return ok - COMILLAS_CURVAS


SIMBOLOS_OK = _simbolos_ok()


def no_permitidos(t):
    return sorted({ch for ch in t if ord(ch) > 127 and ch not in SIMBOLOS_OK
                   and not ('À' <= ch <= 'ÿ' and ch not in '×÷')})


def restos_espanol(t):
    out = []
    limpio = RX_URL.sub(' ', t)
    for f in FRASES_OK:
        limpio = limpio.replace(f, ' ')
    for ch in limpio:
        if ch in SIGNOS_ES:
            out.append(('signo «%s»' % ch, t[:80]))
    for m in RX_PALABRA.finditer(limpio):
        w = m.group(0)
        wl = w.lower()
        if (m.start() > 0 and limpio[m.start() - 1] == '-') or limpio[m.end():m.end() + 1] == '-':
            continue                                             # de-scale, de-lime, re-ice…
        if w.isupper() and len(w) <= 3:
            continue                                             # siglas (Y/N, AM, PM, FOH, MOD…)
        ctx = limpio[max(0, m.start() - 30):m.end() + 30].replace('\n', ' ')
        sin = (wl.replace('á', 'a').replace('é', 'e').replace('í', 'i').replace('ó', 'o').replace('ú', 'u')
               .replace('ñ', 'n'))
        if RX_TILDE.search(w) and wl not in LISTA_BLANCA_TILDE:
            out.append(('tilde/eñe «%s»' % w, ctx))
        elif wl in FUNCIONALES_ES:
            out.append(('funcional ES «%s»' % w, ctx))
        elif sin in OFICIO_ES and wl not in LISTA_BLANCA_TILDE:
            out.append(('oficio ES «%s»' % w, ctx))
    m = RX_COMA_DECIMAL.search(t)
    if m:
        out.append(('coma decimal «%s»' % m.group(0), t[:80]))
    m = RX_HORA_24.search(t)
    if m and not re.search(re.escape(m.group(0)) + r'\s*(?:AM|PM|a\.m\.|p\.m\.)', t):
        out.append(('hora de 24 h «%s»' % m.group(0), t[:80]))
    m = RX_SIGLAS_ES.search(t)
    if m:
        out.append(('sigla ES «%s»' % m.group(0), t[:80]))
    if 'Restaurante Casual' in t or 'Kit de Tareas' in t or any(v in t for v in getattr(mapas, 'VARIANTES_ES', ())):
        out.append(('variante o marca ES', t[:80]))
    return out


def celsius_sueltos(t):
    """D10: todo «°C» de texto va entre paréntesis detrás de su cifra en °F («41 °F (5 °C)»)."""
    malos = []
    for m in re.finditer('°C', t):
        antes = t[:m.start()]
        abre = antes.rfind('(')
        if abre < 0 or ')' in antes[abre:]:
            malos.append(t[max(0, m.start() - 25):m.end() + 5])
            continue
        if '°F' not in antes[abre:] and '°F' not in antes[max(0, abre - 40):abre]:
            malos.append(t[max(0, abre - 25):m.end() + 5])
    return malos


# ==========================================================================
# Utilidades
# ==========================================================================
class Resultado:
    def __init__(self):
        self.gates = OrderedDict()
        self.actual = None

    def gate(self, nombre):
        self.actual = nombre
        self.gates.setdefault(nombre, {'fallos': [], 'datos': OrderedDict()})
        return self

    def fallo(self, msg):
        self.gates[self.actual]['fallos'].append(msg)

    def dato(self, k, v):
        self.gates[self.actual]['datos'][k] = v


_CACHE = {}


def cargar(path, data_only=False):
    k = (path, data_only, os.path.getmtime(path))
    if k not in _CACHE:
        _CACHE[k] = openpyxl.load_workbook(path, data_only=data_only)
    return _CACHE[k]


def libros(carpeta):
    for es in mapas.LIBROS_ES:
        en = mapas.FICHEROS[es]
        yield es, en, X.corto(es), os.path.join(ORIGEN, es), os.path.join(carpeta, en)


def fichero_en(carpeta, corto):
    return [pen for es, en, c, pes, pen in libros(carpeta) if c == corto][0]


def lits(f):
    return [x[1:-1].replace('""', '"') for x in RX_STR.findall(f or '')]


def celdas(sqref):
    out = []
    for rg in str(sqref).split():
        out.extend(A.rango(rg))
    return out


def es_verde(c):
    return (c.fill is not None and c.fill.fgColor is not None and isinstance(c.fill.fgColor.rgb, str)
            and c.fill.fgColor.rgb.upper().endswith(A.VERDE))


def celdas_reales(ws):
    for c in list(ws._cells.values()):
        if c.__class__.__name__ != 'MergedCell':
            yield c


def _pagina(ws):
    ps = ws.page_setup
    fit = ws.sheet_properties.pageSetUpPr.fitToPage if ws.sheet_properties.pageSetUpPr else None
    return (ps.orientation, ps.fitToWidth, ps.fitToHeight, fit, ws.print_title_rows, ws.print_title_cols)


def _area(ws):
    pa = ws.print_area
    return pa.split('!')[-1] if pa else None


def _proteccion(ws):
    p = ws.protection
    return (bool(p.sheet), bool(p.password), p.formatColumns, p.formatRows, p.insertRows, p.deleteRows, p.sort,
            p.autoFilter, p.selectLockedCells, p.selectUnlockedCells, p.formatCells, p.insertColumns,
            p.deleteColumns)


def _cf_firma(cf, regla, trans):
    dxf = regla.dxf
    color = None
    if dxf is not None:
        color = (str(dxf.fill.fgColor.rgb) if dxf.fill is not None and dxf.fill.fgColor is not None else None,
                 str(dxf.font.color.rgb) if dxf.font is not None and dxf.font.color is not None else None)
    return (str(cf.sqref), regla.type, regla.operator, regla.priority, bool(regla.stopIfTrue),
            tuple(trans(x) for x in (regla.formula or [])), regla.text, color)


def _cf_trans(corto, h_es, cf, T):
    """Fórmula EN esperada de una regla de CF: la de A.CF_POR_CELDA si la regla es la excepción declarada, si no ES
    transformado (HOJAS + CLAVES). Una excepción con otra fórmula ES no se admite: sale marcada y G1 falla."""
    def f(x):
        try:
            return A.cf_formula_en(corto, h_es, str(cf.sqref), x, T)[0]
        except A.Aborta as e:
            return '<<%s>>' % e
    return f


def num(v):
    return isinstance(v, (int, float)) and not isinstance(v, bool)


def tareas(ws):
    """Filas de tarea (como el censo): nº entero en A y texto en B, desde la fila 5."""
    out = set()
    for (r, c), cel in list(ws._cells.items()):
        if c == 1 and r >= 5 and isinstance(cel.value, int) and not isinstance(cel.value, bool):
            b = ws._cells.get((r, 2))
            if b is not None and isinstance(b.value, str) and b.data_type != 'f':
                out.add(r)
    return out


def nombres_definidos_crudos(pen):
    """workbook.xml → {(nombre, localSheetId): valor} de _xlnm.Print_Area / _xlnm.Print_Titles (sin desescapar
    más que las entidades XML), y la lista de hojas en orden."""
    with zipfile.ZipFile(pen) as z:
        wbx = z.read('xl/workbook.xml').decode('utf-8')
    hojas = [html.unescape(m.group(1)) for m in re.finditer(r'<sheet\b[^>]*\bname="([^"]*)"', wbx)]
    dn = {}
    for m in re.finditer(r'<definedName\b([^>]*)>(.*?)</definedName>', wbx, re.S):
        nombre = re.search(r'\bname="([^"]*)"', m.group(1)).group(1)
        lid = re.search(r'\blocalSheetId="(\d+)"', m.group(1))
        dn[(nombre, int(lid.group(1)) if lid else None)] = html.unescape(m.group(2))
    return hojas, dn


# ==========================================================================
# G1 — paridad estructural
# ==========================================================================
def gateG1(R, carpeta, datos):
    R.gate('G1 paridad')
    censo = datos['censo']
    tot = Counter()
    for es, en, corto, pes, pen in libros(carpeta):
        if not os.path.isfile(pen):
            R.fallo('%s: no existe' % en)
            continue
        ce = censo['libros'][es]
        wes, wen = cargar(pes), cargar(pen)
        esperado = [mapas.HOJAS[h] for h in ce['hojas']]
        if wen.sheetnames != esperado or wes.sheetnames != ce['hojas']:
            R.fallo('%s: hojas %r ≠ %r' % (en, wen.sheetnames, esperado))
            continue
        tot['hojas'] += len(esperado)
        mapa = OrderedDict((h, mapas.HOJAS[h]) for h in wes.sheetnames)
        T = A.Transformador(mapa)
        hojas_xml, dn = nombres_definidos_crudos(pen)
        for i, h_es in enumerate(wes.sheetnames):
            ws_es, ws_en = wes[h_es], wen[mapa[h_es]]
            hd = ce['hojas_detalle'][h_es]
            lugar = '%s:%s' % (en, ws_en.title)
            # fórmulas
            f_es = {c.coordinate: c.value for c in ws_es._cells.values() if c.data_type == 'f'}
            f_en = {c.coordinate: c.value for c in ws_en._cells.values() if c.data_type == 'f'}
            if len(f_es) != hd['celdas_formula']:
                R.fallo('%s: el ES ya no cuadra con el censo (%d ≠ %d fórmulas)' % (lugar, len(f_es), hd['celdas_formula']))
            if set(f_en) != set(f_es):
                R.fallo('%s: celdas con fórmula distintas: solo ES %s · solo EN %s'
                        % (lugar, sorted(set(f_es) - set(f_en))[:5], sorted(set(f_en) - set(f_es))[:5]))
            con_hoja = 0
            for k, fe in f_es.items():
                fn = f_en.get(k)
                if fn is None:
                    continue
                tot['formulas'] += 1
                try:
                    esperado_f, por_celda = A.formula_en(fe, corto, h_es, k, T)
                except A.Aborta as e:
                    esperado_f, por_celda = '<<%s>>' % e, False
                tot['formulas_por_celda'] += por_celda
                if esperado_f != fn:
                    R.fallo('%s!%s: fórmula EN %r ≠ esperada %r' % (lugar, k, fn[:80], esperado_f[:80]))
                if X.refs_hoja(fn):
                    con_hoja += 1
                    for h in X.refs_hoja(fn):
                        if h not in wen.sheetnames:
                            R.fallo('%s!%s cita la hoja inexistente %r' % (lugar, k, h))
            tot['formulas_con_hoja'] += con_hoja
            con_hoja_esp = hd['formulas_con_hoja'] + CON_HOJA_EXTRA.get((corto, h_es), 0)
            if con_hoja != con_hoja_esp:
                R.fallo('%s: %d fórmulas con referencia a hoja ≠ censo %d' % (lugar, con_hoja, con_hoja_esp))
            # DV
            cen_dv = {d['sqref']: d for d in hd['dv']}
            dves = sorted(ws_es.data_validations.dataValidation, key=lambda d: str(d.sqref))
            dven = sorted(ws_en.data_validations.dataValidation, key=lambda d: str(d.sqref))
            if len(dves) != len(dven) or len(dves) != len(hd['dv']):
                R.fallo('%s: DV ES %d, EN %d, censo %d' % (lugar, len(dves), len(dven), len(hd['dv'])))
            for de, dn_ in zip(dves, dven):
                tot['dv'] += 1
                tot['dv_celdas'] += len(celdas(dn_.sqref))
                firma = lambda d: (str(d.sqref), bool(d.allow_blank), bool(d.showDropDown),  # noqa: E731
                                   bool(d.showErrorMessage), bool(d.showInputMessage), d.errorStyle)
                if firma(de) != firma(dn_):
                    R.fallo('%s: DV %s cambia (%r → %r)' % (lugar, de.sqref, firma(de), firma(dn_)))
                cen = cen_dv.get(str(de.sqref))
                if cen is None or len(celdas(dn_.sqref)) != cen['n_celdas']:
                    R.fallo('%s: DV %s fuera del censo o con otro nº de celdas' % (lugar, dn_.sqref))
                    continue
                for attr in ('error', 'errorTitle', 'prompt', 'promptTitle'):
                    if bool(getattr(de, attr)) != bool(getattr(dn_, attr)):
                        R.fallo('%s: DV %s pierde/gana %s' % (lugar, dn_.sqref, attr))
                f1 = de.formula1 or ''
                try:
                    if de.type == 'list' and f1.startswith('"'):
                        esp = (de.type, de.operator, A.dv_lista_en(f1.strip('"').split(','), lugar), de.formula2)
                    elif de.type == 'custom' and f1:
                        esp = (de.type, de.operator, T.formula(f1), de.formula2)
                    else:
                        esp = (de.type, de.operator, T.refs(f1) if f1 else de.formula1,
                               T.refs(de.formula2) if de.formula2 else de.formula2)
                except A.Aborta as e:
                    esp = ('<<%s>>' % e,)
                got = (dn_.type, dn_.operator, dn_.formula1, dn_.formula2)
                if got != esp:
                    R.fallo('%s: DV %s %r ≠ esperada %r' % (lugar, dn_.sqref, got, esp))
            # CF
            cfe = sorted((_cf_firma(cf, r, _cf_trans(corto, h_es, cf, T)) for cf in ws_es.conditional_formatting
                          for r in cf.rules), key=repr)
            if A.CF_POR_CELDA:
                tot['cf_por_celda'] += sum(1 for cf in ws_es.conditional_formatting
                                           if (corto, h_es, str(cf.sqref)) in A.CF_POR_CELDA)
            cfn = sorted((_cf_firma(cf, r, lambda x: x) for cf in ws_en.conditional_formatting for r in cf.rules), key=repr)
            tot['cf'] += len(cfn)
            if len(cfn) != len(hd['cf']):
                R.fallo('%s: %d CF EN ≠ censo %d' % (lugar, len(cfn), len(hd['cf'])))
            if cfe != cfn:
                dif = [(a, b) for a, b in zip(cfe, cfn) if a != b] or [(cfe[:1], cfn[:1])]
                R.fallo('%s: formato condicional distinto (ES transformado → EN): %r' % (lugar, dif[0]))
            # merges, protección, paneles, impresión, anchos, columnas ocultas
            mg = sorted(map(str, ws_en.merged_cells.ranges))
            if mg != sorted(hd['merges']):
                R.fallo('%s: combinaciones distintas' % lugar)
            tot['merges'] += len(mg)
            if _proteccion(ws_en) != _proteccion(ws_es) or ws_en.protection.password \
                    or bool(ws_en.protection.sheet) != hd['proteccion']['hoja']:
                R.fallo('%s: protección distinta del ES %r ≠ %r' % (lugar, _proteccion(ws_en), _proteccion(ws_es)))
            tot['hojas_protegidas'] += 1 if ws_en.protection.sheet else 0
            if ws_en.freeze_panes != hd['paneles']:
                R.fallo('%s: paneles %s ≠ %s' % (lugar, ws_en.freeze_panes, hd['paneles']))
            tot['paneles'] += 1 if ws_en.freeze_panes else 0
            if _pagina(ws_es) != _pagina(ws_en):
                R.fallo('%s: ajuste o títulos de impresión %r ≠ %r' % (lugar, _pagina(ws_es), _pagina(ws_en)))
            tot['titulos_impresion'] += 1 if ws_en.print_title_rows else 0
            ae = hd['area_impresion'].split('!')[-1] if hd['area_impresion'] else None
            if _area(ws_en) != ae:
                R.fallo('%s: área de impresión %s ≠ %s' % (lugar, _area(ws_en), ae))
            tot['areas_impresion'] += 1 if ae else 0
            # definedName crudo: la hoja citada es la EN, bien entrecomillada ('Valentine''s Day'!)
            if hojas_xml[i] != ws_en.title:
                R.fallo('%s: workbook.xml nombra la hoja %d como %r' % (lugar, i, hojas_xml[i]))
            refs_ok = {A.ref_hoja(ws_en.title) + '!', "'" + ws_en.title.replace("'", "''") + "'!"}
            for nombre, rango_es in (('_xlnm.Print_Area', ae),
                                     ('_xlnm.Print_Titles', hd['titulos_impresion']['filas'])):
                v = dn.get((nombre, i))
                if rango_es is None:
                    continue
                tot['nombres_definidos'] += 1
                partes = [p for p in (v or '').split(',') if p]
                pref = [next((r for r in refs_ok if p.startswith(r)), None) for p in partes]
                if not partes or None in pref:
                    R.fallo('%s: %s crudo %r no cita la hoja %r' % (lugar, nombre, v, ws_en.title))
                elif nombre == '_xlnm.Print_Area' and partes[0][len(pref[0]):] != rango_es:
                    R.fallo('%s: %s crudo %r ≠ %s' % (lugar, nombre, v, rango_es))
                if "'" in ws_en.title:
                    tot['nombres_con_apostrofo'] += 1
            anchos = OrderedDict((k, d.width) for k, d in sorted(ws_en.column_dimensions.items()) if d.width)
            if anchos != OrderedDict(hd['anchos']):
                R.fallo('%s: anchos %r ≠ censo %r' % (lugar, dict(anchos), hd['anchos']))
            tot['anchos'] += len(anchos)
            ocultas = sorted(k for k, d in ws_en.column_dimensions.items() if d.hidden)
            if ocultas != hd['columnas_ocultas']:
                R.fallo('%s: columnas ocultas %s ≠ %s' % (lugar, ocultas, hd['columnas_ocultas']))
            tot['columnas_ocultas'] += len(ocultas)
            # verdes y desbloqueadas, celda a celda
            ue = {c.coordinate for c in celdas_reales(ws_es) if c.protection.locked is False}
            un = {c.coordinate for c in celdas_reales(ws_en) if c.protection.locked is False}
            ve = {c.coordinate for c in celdas_reales(ws_es) if es_verde(c)}
            vn = {c.coordinate for c in celdas_reales(ws_en) if es_verde(c)}
            if un != ue:
                R.fallo('%s: desbloqueadas distintas (solo ES %s, solo EN %s)' % (lugar, sorted(ue - un)[:5],
                                                                                   sorted(un - ue)[:5]))
            if vn != ve:
                R.fallo('%s: verdes distintas (solo ES %s, solo EN %s)' % (lugar, sorted(ve - vn)[:5], sorted(vn - ve)[:5]))
            if ws_en.protection.sheet and vn - un:              # sin protección, el bloqueo no hace nada (catering)
                R.fallo('%s: celdas verdes bloqueadas %s' % (lugar, sorted(vn - un)[:5]))
            tot['verdes'] += len(vn)
            tot['desbloqueadas'] += len(un)
            # tareas: mismas filas (nº en A y texto en B)
            if h_es != 'Instrucciones':
                te_, tn_ = tareas(ws_es), tareas(ws_en)
                if tn_ != te_ or len(tn_) != hd['tareas']:
                    R.fallo('%s: tareas en filas distintas (solo ES %s, solo EN %s; censo %d)'
                            % (lugar, sorted(te_ - tn_)[:5], sorted(tn_ - te_)[:5], hd['tareas']))
                tot['tareas'] += len(tn_)
        with zipfile.ZipFile(pen) as z:
            n = z.namelist()
        g = [x for x in n if x.startswith('xl/charts/') or x.startswith('xl/media/') or x.startswith('xl/drawings/')]
        if g:
            R.fallo('%s: gráficos/imágenes %s' % (en, g[:3]))
    T0 = censo['totales']
    esperados = OrderedDict([
        ('hojas', T0['hojas']), ('formulas', T0['celdas_formula']),
        ('formulas_por_celda', len(mapas.FORMULAS_POR_CELDA)),
        ('formulas_con_hoja', T0['formulas_con_hoja'] + sum(CON_HOJA_EXTRA.values())),
        ('dv', T0['dv']), ('dv_celdas', T0['dv_celdas']), ('cf', T0['cf']), ('merges', T0['merges']),
        ('hojas_protegidas', T0['hojas_protegidas']), ('areas_impresion', T0['areas_impresion']),
        ('titulos_impresion', T0['titulos_impresion']), ('paneles', T0['paneles']),
        ('columnas_ocultas', T0['columnas_ocultas']), ('verdes', T0['verdes']),
        ('desbloqueadas', T0['desbloqueadas']), ('tareas', T0['tareas']),
        ('nombres_definidos', T0['areas_impresion'] + T0['titulos_impresion']),
    ])
    for k, v in esperados.items():
        if tot[k] != v:
            R.fallo('total %s = %d ≠ esperado %d (censo)' % (k, tot[k], v))
    con_apostrofo = [mapas.HOJAS[h] for hs in mapas.HOJAS_ES_POR_LIBRO.values() for h in hs if "'" in mapas.HOJAS[h]]
    if con_apostrofo and not tot['nombres_con_apostrofo']:
        R.fallo('ningún definedName de %s comprobado' % ' / '.join('«%s»' % h for h in con_apostrofo))
    for k, v in sorted(tot.items()):
        R.dato(k, v)


# ==========================================================================
# G2 — restos y caracteres
# ==========================================================================
def piezas_de_texto(pen):
    wb = cargar(pen)
    for ws in wb.worksheets:
        yield 'hoja', ws.title, ws.title, None
        for c in ws._cells.values():
            if isinstance(c.value, str) and c.data_type != 'f':
                yield 'valor', '%s!%s' % (ws.title, c.coordinate), c.value, (ws.title, c.coordinate)
            if c.data_type == 'f':
                for x in lits(c.value):
                    yield 'literal', '%s!%s' % (ws.title, c.coordinate), x, None
            if isinstance(c.number_format, str) and c.number_format != 'General':
                yield 'formato', '%s!%s' % (ws.title, c.coordinate), c.number_format, None
        for dv in ws.data_validations.dataValidation:
            for attr in ('error', 'errorTitle', 'prompt', 'promptTitle'):
                v = getattr(dv, attr)
                if v:
                    yield 'dv', '%s DV %s %s' % (ws.title, dv.sqref, attr), v, None
            f1 = dv.formula1 or ''
            if f1.startswith('"'):
                for x in f1.strip('"').split(','):
                    if x:
                        yield 'dv-item', '%s DV %s' % (ws.title, dv.sqref), x, None
            elif dv.type == 'custom':
                for x in lits(f1):
                    yield 'literal', '%s DV %s' % (ws.title, dv.sqref), x, None
        for cf in ws.conditional_formatting:
            for regla in cf.rules:
                for f in regla.formula or []:
                    for x in lits(f):
                        yield 'literal', '%s CF %s' % (ws.title, cf.sqref), x, None
                if regla.text:
                    yield 'literal', '%s CF %s text' % (ws.title, cf.sqref), regla.text, None
        for parte in ('oddHeader', 'oddFooter', 'evenHeader', 'evenFooter', 'firstHeader', 'firstFooter'):
            hf = getattr(ws, parte)
            for pos in ('left', 'center', 'right'):
                t = getattr(hf, pos).text
                if t:
                    yield 'cabecera-pie', '%s %s.%s' % (ws.title, parte, pos), t, None
    p = wb.properties
    for k in ('title', 'subject', 'creator', 'keywords', 'description', 'category', 'lastModifiedBy'):
        v = getattr(p, k)
        if v:
            yield 'docprops', k, v, None
    with zipfile.ZipFile(pen) as z:
        estilos = z.read('xl/styles.xml').decode('utf-8')
    for m in re.finditer(r'<numFmt [^>]*formatCode="([^"]*)"', estilos):
        yield 'styles.xml', 'numFmt', html.unescape(m.group(1)), None


def zona_dolar():
    """(libro, hoja EN, celda) donde «$»/«¢» son legítimos: los rótulos de POR_CELDA que los llevan (tabla de
    billetes D12) y las celdas cuyo ADAPTAR los pide (nota de la tabla en 09 Instructions)."""
    z = set()
    for (c, h, coord), v in mapas.POR_CELDA.items():
        if '$' in v or '¢' in v:
            z.add((c, mapas.HOJAS[h], coord))
    for (c, h, coord), v in mapas.ADAPTAR.items():
        if '$' in v or '¢' in v:
            z.add((c, mapas.HOJAS[h], coord))
    return z


RX_PROHIBIDOS = re.compile(mapas.TERMINOS_PROHIBIDOS) if mapas.TERMINOS_PROHIBIDOS else None


def gateG2(R, carpeta, datos):
    R.gate('G2 restos')
    n = Counter()
    zona = zona_dolar()
    censo = datos['censo']
    for es, en, corto, pes, pen in libros(carpeta):
        for tipo, lugar, t, pos in piezas_de_texto(pen):
            n[tipo] += 1
            donde = '%s %s' % (en, lugar)
            if '€' in t or re.search(r'\bEUR\b|\beuros?\b', t):
                R.fallo('%s: moneda € en %s: %r' % (donde, tipo, t[:70]))
            malos = no_permitidos(t)
            if malos:
                R.fallo('%s: carácter fuera de la lista blanca %r en %s: %r' % (donde, ''.join(malos), tipo, t[:60]))
            if tipo in ('formato', 'styles.xml'):
                if re.search(r'dd/mm|yyyy-mm-dd|€', t, re.I):
                    R.fallo('%s: formato %r' % (donde, t))
                continue
            if tipo == 'literal':
                continue                                         # ✓, ?*, N/A, >0: los vigila G1 (= ES por CLAVES)
            if RX_PROHIBIDOS is not None:                        # vertical (mapas.TERMINOS_PROHIBIDOS): toda pieza EN
                n['terminos_revisados'] += 1
                m = RX_PROHIBIDOS.search(t)
                if m:
                    R.fallo('%s: término prohibido %r en %s · …%s…' % (donde, m.group(0), tipo,
                                                                       t[max(0, m.start() - 30):m.end() + 20]))
            if ('$' in t or '¢' in t):
                n['con_dolar'] += 1
                if (corto,) + (pos or ('', '')) not in zona:
                    R.fallo('%s: «$»/«¢» fuera de la tabla de billetes (D12) en %s: %r' % (donde, tipo, t[:70]))
            if any(ch in t for ch in COMILLAS_CURVAS):
                R.fallo('%s: comillas curvas en %s: %r' % (donde, tipo, t[:70]))
            if A.RX_RELLENO.search(t):
                n['relleno'] += 1
                if not datos['parcial']:
                    R.fallo('%s: marca de relleno %r (falta la traducción)' % (donde, A.RX_RELLENO.search(t).group(0)))
            if tipo == 'valor' and mapas.RX_HORA_CELDA.match(t):
                R.fallo('%s: hora de 24 h %r sin pasar a reloj de 12 h (D9)' % (donde, t))
            m = RX_FICHERO_ES.search(t)
            if m:
                R.fallo('%s: nombre de fichero ES %r (D5)' % (donde, m.group(0)))
            m = RX_ID_INTERNO.search(t)
            if m:
                R.fallo('%s: ID interno %r en %s: …%s…' % (donde, m.group(0), tipo, t[max(0, m.start() - 30):m.end() + 10]))
            m = RX_CITA_ANIDADA.search(t)
            if m:
                R.fallo('%s: paréntesis anidados %r en %s: …%s…' % (donde, m.group(0), tipo, t[max(0, m.start() - 40):m.end() + 30]))
            if tipo == 'dv':
                attr = lugar.rsplit(' ', 1)[-1]
                if len(t) > LIMITE_DV[attr]:
                    R.fallo('%s: DV %s de %d caracteres > %d (límite de Excel)' % (donde, attr, len(t), LIMITE_DV[attr]))
                n['dv_' + attr] += 1
            for motivo, ctx in restos_espanol(t):
                R.fallo('%s: %s · …%s…' % (donde, motivo, ctx))
            for s in celsius_sueltos(t):
                R.fallo('%s: °C sin su °F delante · …%s…' % (donde, s))
            if '°C' in t:
                n['textos_con_celsius_ok'] += 1
        # formatos con € del censo: reescritos
        ce = censo['libros'][es]
        wen = cargar(pen)
        for h_es, hd in ce['hojas_detalle'].items():
            ws = wen[mapas.HOJAS[h_es]]
            for coord in hd['formatos_euro']:
                n['formatos_euro_reescritos'] += 1
                f = ws[coord].number_format
                if '€' in f or f != A.FORMATOS.get(censo_formato(pes, h_es, coord), f):
                    R.fallo('%s %s!%s: formato %r (D13)' % (en, ws.title, coord, f))
    if n['formatos_euro_reescritos'] != censo['totales']['formatos_euro']:
        R.fallo('formatos con € revisados %d ≠ censo %d' % (n['formatos_euro_reescritos'], censo['totales']['formatos_euro']))
    for k, v in sorted(n.items()):
        R.dato(k, v)


def censo_formato(pes, h_es, coord):
    return cargar(pes)[h_es][coord].number_format


# ==========================================================================
# G3 — formato
# ==========================================================================
def gateG3(R, carpeta, datos, mes):
    R.gate('G3 formato')
    censo = datos['censo']
    n = Counter()
    rx_version = re.compile(r'^Version %s · %s \d{4} · %s · info@aichef\.pro$'
                            % (re.escape(A.VERSION), mes, re.escape(A.DESCRIPTION)))
    for es, en, corto, pes, pen in libros(carpeta):
        ce = censo['libros'][es]
        wb = cargar(pen)
        wes = cargar(pes)
        marcas_es = sum(1 for ws in wes.worksheets for c in ws._cells.values() if c.value == mapas.MARCA_ES)
        marcas_en = sum(1 for ws in wb.worksheets for c in ws._cells.values() if c.value == mapas.MARCA_EN)
        if marcas_en != marcas_es or (marcas_es == 0 and corto not in SIN_MARCA):
            R.fallo('%s: %d líneas de marca D19 ≠ %d del ES' % (en, marcas_en, marcas_es))
        n['marcas'] += marcas_en
        for h_es, hd in ce['hojas_detalle'].items():
            ws = wb[mapas.HOJAS[h_es]]
            ws_es = wes[h_es]
            n['hojas'] += 1
            if ws.page_setup.paperSize not in (1, '1'):
                R.fallo('%s:%s: paperSize %s ≠ 1 (Letter)' % (en, ws.title, ws.page_setup.paperSize))
            for k in hd['cabeceras_pies']:
                obj, pos = k.split('.')
                v = getattr(getattr(ws, obj), pos).text
                if v != A.PIE_EN:
                    R.fallo('%s:%s: %s = %r ≠ pie D19' % (en, ws.title, k, v))
                else:
                    n['pies'] += 1
            # ajuste de texto y alto de fila (la misma cuenta que aplicar_en.capa_ajuste)
            anclas, cubiertas = A.mapa_merges(ws)
            for c in list(ws._cells.values()):
                necesita, alto, desborde = A.objetivo_ajuste(ws, ws_es, c, anclas, cubiertas)
                if desborde:
                    R.fallo('%s:%s!%s: el EN no cabe ni con ajuste (acortarlo): %r' % (en, ws.title, c.coordinate,
                                                                                        c.value[:60]))
                if necesita:
                    n['ajuste_necesario'] += 1
                    if not (c.alignment and c.alignment.wrap_text):
                        R.fallo('%s:%s!%s: el EN desborda y la celda no tiene ajuste de texto: %r'
                                % (en, ws.title, c.coordinate, c.value[:60]))
                if alto is not None:
                    n['alto_necesario'] += 1
                    dim = ws.row_dimensions.get(c.row)
                    actual = dim.height if dim is not None else None
                    if (actual or A.ALTO_DEFECTO) + 1e-9 < alto:
                        R.fallo('%s:%s!%s: fila %d de alto %s < %s para el EN (ajuste de texto)'
                                % (en, ws.title, c.coordinate, c.row, actual, alto))
            # ningún alto de fila baja respecto al ES
            for r, d in ws_es.row_dimensions.items():
                if d.height:
                    dn_ = ws.row_dimensions.get(r)
                    if dn_ is None or not dn_.height or dn_.height + 1e-9 < d.height:
                        R.fallo('%s:%s: la fila %d baja de %s a %s' % (en, ws.title, r, d.height,
                                                                     dn_.height if dn_ is not None else None))
        ins = wb[A.HOJA_INSTR_EN]
        vers = [c.value for c in ins._cells.values() if isinstance(c.value, str) and c.value.startswith('Version ')]
        if len(vers) != 1 or not rx_version.match(vers[0]):
            R.fallo('%s: línea de versión D19 %r' % (en, vers))
        else:
            n['versiones'] += 1
        p = wb.properties
        esp = {'title': mapas.TITULOS[en], 'subject': A.SUBJECT, 'keywords': A.KEYWORDS, 'description': A.DESCRIPTION,
               'category': A.CATEGORY, 'creator': A.CREATOR}
        for k, v in esp.items():
            if getattr(p, k) != v:
                R.fallo('%s: docProps %s %r ≠ %r' % (en, k, getattr(p, k), v))
        if ins['B2'].value != mapas.titulo_interior(es):
            R.fallo('%s: Instructions!B2 %r ≠ título interior D5 %r' % (en, ins['B2'].value, mapas.titulo_interior(es)))
    T0 = censo['totales']
    n_pies = sum(len(hd['cabeceras_pies']) for i in censo['libros'].values() for hd in i['hojas_detalle'].values())
    if n['hojas'] != T0['hojas'] or n['pies'] != n_pies:
        R.fallo('hojas %d / pies %d ≠ censo %d / %d' % (n['hojas'], n['pies'], T0['hojas'], n_pies))
    if n['versiones'] != len(mapas.FICHEROS):
        R.fallo('versiones %d ≠ %d' % (n['versiones'], len(mapas.FICHEROS)))
    for k, v in sorted(n.items()):
        R.dato(k, v)


# ==========================================================================
# G4 — mapas (por aparición)
# ==========================================================================
def gateG4(R, carpeta, datos, mes):
    R.gate('G4 mapas')
    n = Counter()
    te = datos['te']
    tx = datos['tx']

    def norm(v):
        for a, b in A.COMILLAS.items():
            v = v.replace(a, b)
        return v
    wbs = {c: cargar(fichero_en(carpeta, c)) for c in (X.corto(f) for f in mapas.LIBROS_ES)}
    for cad in te['cadenas']:
        ident, es = cad['id'], cad['es']
        en_tx = tx.get(ident)
        if en_tx is None and datos['parcial'] and cad['grupo'] != 'GM':
            en_tx = A.RELLENO % ident
        for d in cad['donde']:
            t, tr = d['t'], d['tr']
            if t in ('hoja', 'docprops', 'cabecera-pie'):
                continue                                         # G1 (hojas) y G3 (docProps, pies)
            ws = wbs[d['f']][mapas.HOJAS[d['h']]]
            if t == 'celda':
                got = ws[d['c']].value
                if tr in ('traducir', 'localizar'):
                    esp = norm(en_tx) if en_tx is not None else None
                elif tr == 'mapa-glosario':
                    esp = mapas.GLOSARIO[es]
                elif tr == 'mapa-hoja':
                    esp = mapas.HOJAS[es]
                elif tr == 'mapa-clave':
                    esp = CLAVES[es]
                else:
                    esp = A.regenerado_esperado(d['f'], d['h'], d['c'], es, mes)
                n['celda_' + tr] += 1
                if esp is None or got != esp:
                    R.fallo('%s %s!%s (%s %s): %r ≠ %r' % (d['f'], ws.title, d['c'], ident, tr,
                                                          (got or '')[:70] if isinstance(got, str) else got,
                                                          (esp or '')[:70] if not isinstance(esp, (int, float))
                                                          else esp))
            elif t in A.CAMPO_DV:
                attr = A.CAMPO_DV[t]
                dvs = [x for x in ws.data_validations.dataValidation if str(x.sqref) == d['c']]
                got = getattr(dvs[0], attr) if len(dvs) == 1 else None
                esp = A.DV_MENSAJES_EN.get((d['f'], d['h'], d['c'], attr), norm(en_tx) if en_tx else None)
                n['dv_mensajes'] += 1
                if got != esp:
                    R.fallo('%s %s DV %s %s (%s): %r ≠ %r' % (d['f'], ws.title, d['c'], attr, ident, got, esp))
        # citas de pestañas/rótulos y ficheros con su EN exacto (SPEC §7.2)
        if cad['tratamiento'] in ('traducir', 'localizar') and en_tx and not A.RX_RELLENO.match(en_tx):
            for c in cad.get('cita_rotulos', []):
                n['citas_rotulos'] += 1
                if c['en'] and c['en'] not in en_tx:
                    R.fallo('%s: cita «%s» sin su EN exacto %r · %r' % (ident, c['es'], c['en'], en_tx[:90]))
            for c in cad.get('cita_ficheros', []):
                n['citas_ficheros'] += 1
                if c['en'] not in en_tx:
                    R.fallo('%s: fichero %s sin su nombre EN %s · %r' % (ident, c['es'], c['en'], en_tx[:90]))
            if es.startswith('Diseñado por John Guerrero'):
                n['firmas'] += 1
                if en_tx != FIRMA_EN:
                    R.fallo('%s: firma de autor %r ≠ SPEC §7.2 %r' % (ident, en_tx, FIRMA_EN))
    # horas de celda (D9): cada «HH:MM» del ES → mapas.hora_us en la misma celda
    for es_f, en_f, corto, pes, pen in libros(carpeta):
        wes, wen = cargar(pes), cargar(pen)
        for h_es in wes.sheetnames:
            ws_es, ws_en = wes[h_es], wen[mapas.HOJAS[h_es]]
            for c in list(ws_es._cells.values()):
                if isinstance(c.value, str) and c.data_type != 'f' and mapas.RX_HORA_CELDA.match(c.value):
                    n['horas'] += 1
                    got = ws_en[c.coordinate].value
                    esp = mapas.POR_CELDA.get((corto, h_es, c.coordinate)) or mapas.hora_us(c.value)
                    if got != esp:                       # POR_CELDA manda (cafe ronda 1: horas alineadas)
                        R.fallo('%s %s!%s: hora %r ≠ hora_us(%r) = %r' % (corto, ws_en.title, c.coordinate, got,
                                                                          c.value, esp))
    if n['horas'] != datos['censo']['totales']['horas_24h']:
        R.fallo('horas de celda %d ≠ censo %d' % (n['horas'], datos['censo']['totales']['horas_24h']))
    # números fijados (mapas.VALORES_EN; vacío en el base y el bar): el ES trae el suyo y el EN el número decidido
    wes_v = {}
    for (c, h_es, coord), (v_es, v_en) in mapas.VALORES_EN.items():
        n['valores_fijados'] += 1
        if c not in wes_v:
            wes_v[c] = cargar(next(p for _e, _en, corto, p, _pen in libros(carpeta) if corto == c))
        got_es = wes_v[c][h_es][coord].value
        got = wbs[c][mapas.HOJAS[h_es]][coord].value
        if got_es != v_es or got != v_en or isinstance(got, bool):
            R.fallo('%s %s!%s: valor fijado %r (ES %r) ≠ %r (ES %r)' % (c, mapas.HOJAS[h_es], coord, got, got_es,
                                                                       v_en, v_es))
    if not n['firmas'] and not datos['parcial']:
        R.fallo('firma de autor no comprobada (¿cadena ausente?)')
    for k, v in sorted(n.items()):
        R.dato(k, v)


# ==========================================================================
# G5 — caja en dólares (D12) y caso de SPEC §5 en una copia
# ==========================================================================
def recalcular(p):
    r = subprocess.run([sys.executable, '-W', 'ignore', A.INJECT, p], stdout=subprocess.PIPE,
                       stderr=subprocess.PIPE, universal_newlines=True)
    if r.returncode or 'fallos_pycel=0' not in r.stdout:
        return 'inject_cache: %s %s' % (r.stdout.strip()[-120:], r.stderr[-200:])
    return None


def caso_caja(carpeta):
    """Copia del 09 con el caso de SPEC §5. Devuelve (valores, error)."""
    hc, ha, hr = (mapas.HOJAS['Cierre de Caja'], mapas.HOJAS['Apertura de Caja'], mapas.HOJAS['Registro Mensual'])
    filas = [d[0] for d in mapas.DENOMINACIONES]
    out = OrderedDict()
    tmp = tempfile.mkdtemp(prefix='checklists-caja-')
    try:
        src = fichero_en(carpeta, '09')
        # caso 1: una unidad de cada fila
        p = os.path.join(tmp, '1-' + os.path.basename(src))
        shutil.copy2(src, p)
        wb = openpyxl.load_workbook(p)
        for f in filas:
            wb[hc]['B%d' % f].value = 1
        wb.save(p)
        e = recalcular(p)
        if e:
            return None, e
        wv = openpyxl.load_workbook(p, data_only=True)
        out['total_1_de_cada'] = wv[hc]['C43'].value
        # caso 2: cierre con fondo 200, contado 1,250 (12 × $100 + 1 × $50), tarjetas 3,000, otros 150, Z 4,200
        p = os.path.join(tmp, '2-' + os.path.basename(src))
        shutil.copy2(src, p)
        wb = openpyxl.load_workbook(p)
        caras = {d[0]: d[4] for d in mapas.DENOMINACIONES}
        resto = CAJA_SPEC['contado']
        for f in sorted(filas, key=lambda x: -caras[x]):
            if f in (41, 42) or caras[f] > resto + 1e-9:          # los cartuchos no se usan en el caso
                wb[hc]['B%d' % f].value = 0
                continue
            q = int(resto // caras[f])
            wb[hc]['B%d' % f].value = q
            resto = round(resto - q * caras[f], 2)
        wb[ha]['C23'].value = CAJA_SPEC['fondo']
        wb[hc]['C50'].value = CAJA_SPEC['tarjetas']
        wb[hc]['C51'].value = CAJA_SPEC['otros']
        wb[hc]['C53'].value = CAJA_SPEC['z']
        for col, v in (('B', 'fondo'), ('C', 'contado'), ('D', 'tarjetas'), ('E', 'otros'), ('G', 'z')):
            wb[hr]['%s5' % col].value = CAJA_SPEC[v]
        wb.save(p)
        e = recalcular(p)
        if e:
            return None, e
        wv = openpyxl.load_workbook(p, data_only=True)
        for k, coord in (('contado', 'C47'), ('fondo', 'C48'), ('ventas_efectivo', 'C49'), ('total_cobrado', 'C52'),
                         ('over_short', 'C54')):
            out['cierre_' + k] = wv[hc][coord].value
        out['registro_total_cobrado'] = wv[hr]['F5'].value
        out['registro_over_short'] = wv[hr]['H5'].value
        return out, None
    finally:
        shutil.rmtree(tmp)


def casos_g5(carpeta):
    """Vertical con el 09 propio: cada caso de mapas.CASOS_G5 en una copia → [(caso, celda, valor, esperado)], error."""
    import datetime as _dt
    out = []
    tmp = tempfile.mkdtemp(prefix='checklists-g5-')
    try:
        src = fichero_en(carpeta, '09')
        for i, caso in enumerate(CASOS_G5):
            p = os.path.join(tmp, '%d-%s' % (i, os.path.basename(src)))
            shutil.copy2(src, p)
            wb = openpyxl.load_workbook(p)
            for (h_es, coord), v in caso['entradas'].items():
                wb[mapas.HOJAS[h_es]][coord].value = _dt.datetime(2020, 1, 1) if v == 'FECHA_PASADA' else v
            wb.save(p)
            e = recalcular(p)
            if e:
                return None, '%s: %s' % (caso['nombre'], e)
            wv = openpyxl.load_workbook(p, data_only=True)
            for (h_es, coord), esp in caso['esperado'].items():
                out.append((caso['nombre'], '%s!%s' % (mapas.HOJAS[h_es], coord), wv[mapas.HOJAS[h_es]][coord].value,
                            esp))
        return out, None
    finally:
        shutil.rmtree(tmp)


def gateG5(R, carpeta, datos):
    R.gate('G5 caja')
    hc = mapas.HOJAS[HOJA_CAJA[1]]
    ws = cargar(fichero_en(carpeta, '09'))[hc]
    total = 0
    for f, es, xes, en, xen in mapas.DENOMINACIONES:
        k = (HOJA_CAJA[0], HOJA_CAJA[1], 'C%d' % f)
        if ws['A%d' % f].value != en:
            R.fallo('09 %s!A%d = %r ≠ %r (DENOMINACIONES)' % (hc, f, ws['A%d' % f].value, en))
        if ws['C%d' % f].value != mapas.FORMULAS_POR_CELDA[k][1]:
            R.fallo('09 %s!C%d = %r ≠ %r (caja D12)' % (hc, f, ws['C%d' % f].value, mapas.FORMULAS_POR_CELDA[k][1]))
        total += xen
    R.dato('denominaciones', len(mapas.DENOMINACIONES))
    if abs(total - CAJA_SPEC['total_1_de_cada']) > 1e-6:
        R.fallo('mapas.DENOMINACIONES suma %.2f ≠ SPEC §5 %.2f' % (total, CAJA_SPEC['total_1_de_cada']))
    if CASOS_G5 is not None:
        vals, err = casos_g5(carpeta)
        if err:
            R.fallo('casos del 09: %s' % err)
            return
        R.dato('casos', len(CASOS_G5))
        R.dato('comprobaciones', len(vals))
        for nombre, donde, v, esp in vals:
            ok = (num(v) and abs(v - esp) < 0.005) if not isinstance(esp, str) else v == esp
            if not ok:
                R.fallo('caja 09 «%s» %s = %r ≠ %r' % (nombre, donde, v, esp))
        return
    vals, err = caso_caja(carpeta)
    if err:
        R.fallo('caso de caja: %s' % err)
        return
    for k, v in vals.items():
        R.dato(k, v)
    cobrado = CAJA_SPEC['contado'] - CAJA_SPEC['fondo'] + CAJA_SPEC['tarjetas'] + CAJA_SPEC['otros']
    esperado = OrderedDict([('total_1_de_cada', total), ('cierre_contado', CAJA_SPEC['contado']),
                            ('cierre_fondo', CAJA_SPEC['fondo']),
                            ('cierre_ventas_efectivo', CAJA_SPEC['contado'] - CAJA_SPEC['fondo']),
                            ('cierre_total_cobrado', cobrado), ('cierre_over_short', cobrado - CAJA_SPEC['z']),
                            ('registro_total_cobrado', cobrado), ('registro_over_short', cobrado - CAJA_SPEC['z'])])
    if cobrado != 4200 or cobrado - CAJA_SPEC['z'] != 0:
        R.fallo('el caso de SPEC §5 ya no da 4,200 y 0')
    for k, e in esperado.items():
        v = vals.get(k)
        if not (num(v) and abs(v - e) < 0.005):
            R.fallo('caja: %s = %r ≠ %r (SPEC §5)' % (k, v, e))


# ==========================================================================
# G6 — cálculo
# ==========================================================================
def celdas_sin_v(path):
    fuera = set()
    with zipfile.ZipFile(path) as z:
        wbx = z.read('xl/workbook.xml').decode('utf-8')
        rels = z.read('xl/_rels/workbook.xml.rels').decode('utf-8')
        destino = {}
        for tag in re.findall(r'<Relationship\b[^>]*>', rels):
            i, t = re.search(r'\bId="([^"]+)"', tag), re.search(r'\bTarget="([^"]+)"', tag)
            if i and t:
                destino[i.group(1)] = t.group(1)
        for tag in re.findall(r'<sheet\b[^>]*>', wbx):
            nombre = html.unescape(re.search(r'\bname="([^"]*)"', tag).group(1))
            rid = re.search(r'\br:id="([^"]+)"', tag).group(1)
            t = destino[rid]
            parte = t.lstrip('/') if t.startswith('/') else 'xl/' + t
            xml = z.read(parte).decode('utf-8')
            for m in re.finditer(r'<c\b([^>]*?)(?:/>|>(.*?)</c>)', xml, re.S):
                dentro = m.group(2) or ''
                if '<f' in dentro and not re.search(r'<v>[^<]', dentro):
                    fuera.add((nombre, re.search(r'\br="([A-Z]+\d+)"', m.group(1)).group(1)))
    return fuera


def gateG6(R, carpeta, datos):
    R.gate('G6 cálculo')
    n = Counter()
    for es, en, corto, pes, pen in libros(carpeta):
        wf, wv = cargar(pen), cargar(pen, data_only=True)
        wes_v, wes_f = cargar(pes, data_only=True), cargar(pes)
        sin_v = celdas_sin_v(pen)
        sin_v_es = celdas_sin_v(pes)          # «<v></v>» legítimo: la fórmula da "" (p. ej. IF(B6="","",…) vacío)
        for h_es in wes_f.sheetnames:
            h_en = mapas.HOJAS[h_es]
            ws, ws_v, we_v = wf[h_en], wv[h_en], wes_v[h_es]
            for c in ws._cells.values():
                if c.data_type != 'f':
                    continue
                n['formulas'] += 1
                if (ws.title, c.coordinate) in sin_v and (h_es, c.coordinate) in sin_v_es:
                    n['vacias_como_es'] += 1                     # el ES tampoco trae valor: paridad
                    continue
                if (ws.title, c.coordinate) in sin_v:
                    R.fallo('%s:%s!%s sin caché (<f> sin <v>)' % (en, ws.title, c.coordinate))
                    continue
                v = ws_v[c.coordinate].value
                ve = we_v[c.coordinate].value
                if isinstance(v, str) and v in ERRORES:
                    R.fallo('%s:%s!%s = %s' % (en, ws.title, c.coordinate, v))
                    continue
                n['con_valor' if v is not None else 'vacias'] += 1
                if 'COUNT' in c.value.upper():
                    n['contadores'] += 1
                if (ve is None) != (v is None):
                    R.fallo('%s:%s!%s caché %r (ES %r)' % (en, ws.title, c.coordinate, v, ve))
                elif isinstance(ve, str):
                    n['textos'] += 1
                    if v != CLAVES.get(ve, ve):
                        R.fallo('%s:%s!%s texto calculado %r ≠ ES %r' % (en, ws.title, c.coordinate, v, ve))
                elif num(ve):
                    n['numeros'] += 1
                    if not (num(v) and abs(v - ve) < TOL):
                        R.fallo('%s:%s!%s número %r ≠ ES %r' % (en, ws.title, c.coordinate, v, ve))
    if n['formulas'] != datos['censo']['totales']['celdas_formula']:
        R.fallo('fórmulas %d ≠ censo %d' % (n['formulas'], datos['censo']['totales']['celdas_formula']))
    for k, v in sorted(n.items()):
        R.dato(k, v)


# ==========================================================================
# G7 — reglas (SPEC §3 = mapas.REGLAS_US; §8: cifras °F)
# ==========================================================================
RX_GRADOS_F = re.compile(r'([−-]?\d+(?:\.\d+)?)\s*(?:(?:–|-|to|and)\s*([−-]?\d+(?:\.\d+)?)\s*)?°\s*F')
# familia → (regex, nº de grupos que forman el prefijo). Los rangos y listas («516.5-516.6», «117 y 109») se
# expanden: cada número es una cita.
RX_CITAS = OrderedDict([
    ('food_code', re.compile(r'(?<![\d.])(\d-\d{3}\.\d{2})(?![\d])')),
    ('cfr', re.compile(r'\b(\d{1,2}) CFR (?:Part |§ ?)?(\d+(?:\.\d+)?)((?:\s*(?:[-–,]|and|y|&)\s*\d+(?:\.\d+)?)*)')),
    ('osha', re.compile(r'\b(1910\.\d+)')),
    ('nfpa', re.compile(r'\bNFPA (\d+[A-Z]?)')),
    ('epa', re.compile(r'\bEPA (\d+)')),
    ('usc', re.compile(r'\b(\d{1,2}) U\.S\.C\. (?:§ ?)?(\d+)((?:\s*(?:[-–,]|and|y|&)\s*\d+)*)')),
    ('seccion', re.compile(r'§ ?(\d+(?:[-.]\d+)*)')),
])


def _num_f(s):
    return float(s.replace('−', '-'))


def grados_f(t):
    out = []
    for m in RX_GRADOS_F.finditer(t):
        out.append(_num_f(m.group(1)))
        if m.group(2):
            out.append(_num_f(m.group(2)))
    return out


def citas_de(t):
    out = []
    for fam, rx in RX_CITAS.items():
        for m in rx.finditer(t):
            g = m.groups()
            if fam in ('cfr', 'usc'):
                out.append((fam, '%s %s' % (g[0], g[1])))
                for extra in re.findall(r'\d+(?:\.\d+)?', g[2] or ''):
                    out.append((fam, '%s %s' % (g[0], extra)))
            else:
                out.append((fam, g[0]))
    return out


def permitidas():
    """Citas normativas y cifras °F que el kit puede afirmar: las de mapas.REGLAS_US (fuentes y valores),
    mapas.POR_CELDA y mapas.ADAPTAR (que salen de SPEC §3)."""
    citas, grados = set(), set()
    textos = [v for _k, (v, f) in mapas.REGLAS_US.items()] + [f for _k, (v, f) in mapas.REGLAS_US.items()]
    textos += list(mapas.POR_CELDA.values()) + list(mapas.ADAPTAR.values())
    for t in textos:
        citas.update(citas_de(t))
        grados.update(grados_f(t))
    # «29 CFR 1910.157(e)(2)» deja también la familia osha; las subsecciones (e)(2), (e)(3), (f) van dentro
    return citas, grados


def gateG7(R, carpeta, datos):
    R.gate('G7 reglas')
    n = Counter()
    citas_ok, grados_ok = permitidas()
    if grados_ok != {float(x) for x in GRADOS_F_SPEC}:
        R.fallo('las cifras °F de mapas.py %s ≠ SPEC §8 G7 %s' % (sorted(grados_ok), sorted(GRADOS_F_SPEC)))
    for es, en, corto, pes, pen in libros(carpeta):
        wes = cargar(pes)
        for tipo, lugar, t, pos in piezas_de_texto(pen):
            if tipo not in ('valor', 'dv'):
                continue
            for g in grados_f(t):
                n['grados_f'] += 1
                if g not in grados_ok:
                    R.fallo('%s %s: %g °F fuera de SPEC §3/§8 · %r' % (en, lugar, g, t[:90]))
            for c in citas_de(t):
                n['citas_normativas'] += 1
                if c not in citas_ok:
                    R.fallo('%s %s: cita normativa %s %r fuera de SPEC §3 · %r' % (en, lugar, c[0], c[1], t[:90]))
        # D10: toda celda ES con °C lleva °F en el EN
        ce = datos['censo']['libros'][es]
        wen = cargar(pen)
        for h_es, hd in ce['hojas_detalle'].items():
            for coord in hd['textos_celsius']:
                n['celdas_celsius_es'] += 1
                v = wen[mapas.HOJAS[h_es]][coord].value
                if (ce['corto'], h_es, coord) in getattr(mapas, 'CELSIUS_SIN_F', ()):
                    n['celsius_sin_f_declaradas'] += 1
                    if isinstance(v, str) and ('°C' in v or grados_f(v)):
                        R.fallo('%s %s!%s: declarada sin cifra (mapas.CELSIUS_SIN_F) y lleva temperatura: %r'
                                % (en, mapas.HOJAS[h_es], coord, v[:80]))
                    continue
                if not (isinstance(v, str) and (A.RX_RELLENO.match(v) or '°F' in v)):
                    R.fallo('%s %s!%s: el ES tenía °C y el EN no lleva °F: %r (ES %r)'
                            % (en, mapas.HOJAS[h_es], coord, v, wes[h_es][coord].value[:60]))
    if n['celdas_celsius_es'] != datos['censo']['totales']['textos_celsius']:
        R.fallo('celdas con °C %d ≠ censo %d' % (n['celdas_celsius_es'], datos['censo']['totales']['textos_celsius']))
    R.dato('citas_permitidas', len(citas_ok))
    for k, v in sorted(n.items()):
        R.dato(k, v)


# ==========================================================================
# G8 — censo-entregables y gate-no-latinos
# ==========================================================================
def gateG8(R, carpeta, datos):
    R.gate('G8 censo + no-latinos')
    real = os.path.abspath(carpeta) == os.path.abspath(A.DESTINO_REAL)
    cmds = [[sys.executable, CENSO_SCRIPT, '--only', mapas.SLUG if real else carpeta, '--fail'] + ([] if real else ['--letter']),
            [sys.executable, NO_LATINOS, '--only', mapas.SLUG if real else carpeta]]
    for cmd in cmds:
        r = subprocess.run(cmd, stdout=subprocess.PIPE, stderr=subprocess.STDOUT, universal_newlines=True, cwd=REPO)
        lineas = [x for x in r.stdout.strip().splitlines() if x.strip()]
        nombre = os.path.basename(cmd[1])
        R.dato(nombre, lineas[-1] if lineas else '')
        if r.returncode:
            R.fallo('%s exit %d: %s' % (nombre, r.returncode, ' | '.join(lineas[-6:])))


# ==========================================================================
# Autotest: cada gate tiene que cazar un defecto inyectado en una COPIA
# ==========================================================================
def _libro(corto):
    """Fichero EN del libro `corto` («01»…) por su prefijo: los nombres ES cambian por vertical (heladería: 02-partidas-produccion, 06-eventos-temporada)."""
    return mapas.FICHEROS[next(f for f in mapas.LIBROS_ES if f.startswith(corto + '-'))]


F01, F02, F03, F06 = (_libro(c) for c in ('01', '02', '03', '06'))
F09 = _libro('09')                                         # catering: 09-cobros-…
H_KO = mapas.HOJAS[mapas.HOJAS_ES_POR_LIBRO['01'][1]]      # base: «Kitchen Opening» · bar: la 1.ª checklist del 01
H_CC = mapas.HOJAS[HOJA_CAJA[1]]                           # base: «Cash Drawer Closing» · catering: «Event Settlement»
_HH = getattr(mapas, 'HOJA_HORAS', None)                   # catering: las horas de celda están en el 08
F_HORAS = mapas.FICHEROS[next(f for f in mapas.LIBROS_ES if f.startswith(_HH[0] + '-'))] if _HH else F01
H_HORAS = mapas.HOJAS[_HH[1]] if _HH else H_KO
CITA_MAL = getattr(mapas, 'CITA_AUTOTEST', '(FDA Food Code 3-501.19)')
_HV = getattr(mapas, 'HOJA_VERDE', None)                   # catering: el 01 no protege sus verdes (el 08 sí)
F_VERDE = mapas.FICHEROS[next(f for f in mapas.LIBROS_ES if f.startswith(_HV[0] + '-'))] if _HV else F01
H_VERDE = mapas.HOJAS[_HV[1]] if _HV else H_KO
F_CAJA0 = mapas.DENOMINACIONES[0][0]                       # base: 27 · catering: 30
H_AREA = mapas.HOJAS[mapas.HOJAS_ES_POR_LIBRO['06'][2]]    # base: «Valentine's Day»
H_DIARIO = mapas.HOJAS[mapas.HOJAS_ES_POR_LIBRO['03'][1]]  # «Manager Daily»
RX_HORA_US = re.compile(r'^\d{1,2}:\d{2} [AP]M$')


def _primera(ws, preferida, otras):
    """La celda con el valor `preferido` (el del kit base) o, si la hoja no lo trae, la primera que cumple `otras`."""
    c = next((c for c in ws._cells.values() if c.value == preferida), None)
    return c or next(c for c in ws._cells.values() if isinstance(c.value, str) and otras(c.value))


def _mutar(carpeta, fname, fn):
    p = os.path.join(carpeta, fname)
    wb = openpyxl.load_workbook(p)
    fn(wb)
    wb.save(p)


def _set(ws, coord, v):
    ws[coord].value = v


def _bloquear_verde(wb):
    from openpyxl.styles import Protection
    ws = wb[H_VERDE]
    c = next(c for c in ws._cells.values() if es_verde(c) and c.protection.locked is False)
    c.protection = Protection(locked=True)


def _quitar_tarea(wb):
    ws = wb[H_KO]
    r = min(tareas(ws))
    ws['B%d' % r].value = None


def _area_rota(wb):
    wb[H_AREA].print_area = 'A1:C10'


def _celda_glosario(wb):
    ws = wb[H_KO]
    puestos = {mapas.GLOSARIO[k] for k in mapas.PUESTOS_ES if k in mapas.GLOSARIO}
    c = _primera(ws, mapas.GLOSARIO['Todo el equipo'], lambda v: v in puestos)
    c.value = 'Everybody'


def _hora_mal(wb):
    ws = wb[H_HORAS]
    c = _primera(ws, '7:00 AM', RX_HORA_US.match)
    c.value = c.value.replace(':00 ', ' ').replace(':', '.')


def _hora_24(wb):
    ws = wb[H_HORAS]
    c = _primera(ws, '7:00 AM', RX_HORA_US.match)
    c.value = '07:00' if c.value == '7:00 AM' else '17:00'


def _titulo_largo(wb):
    from openpyxl.styles import Alignment
    ws = wb[H_KO]
    ws['A1'].value = 'Kitchen Opening Checklist — ' + 'every station, every walk-in, every shift, ' * 4
    ws['A1'].alignment = Alignment(wrap_text=False)


def _caja_mal(wb):
    wb[H_CC]['C%d' % F_CAJA0].value = '=B%d*50' % F_CAJA0


def _cache_mal(wb):
    wb[H_CC][getattr(mapas, 'CELDA_CACHE_MAL', 'C50')].value = 100


def _grados_mal(wb):
    ws = wb[mapas.HOJAS['Instrucciones']]
    g = next(x for x in range(38, 0, -1) if x not in GRADOS_F_SPEC)   # base: 38 · bar (38 °F es la cerveza): 37
    ws['B6'].value = (ws['B6'].value or '') + ' Hold cold food at %d °F.' % g


def _cita_mal(wb):
    ws = wb[mapas.HOJAS['Instrucciones']]
    if hasattr(mapas, 'CITA_AUTOTEST'):        # el kit fija su cita no permitida
        ws['B6'].value = (ws['B6'].value or '') + ' ' + CITA_MAL
        return
    ok = permitidas()[0]                       # base y bar: 3-501.19 · cafe (la cita en REGLAS_US, ronda 1 M4): la siguiente
    cod = next(c for c in ('3-501.19', '3-501.18', '3-202.11', '3-401.11', '3-304.11', '4-204.112')
               if not set(citas_de('FDA Food Code ' + c)) <= ok)
    ws['B6'].value = (ws['B6'].value or '') + ' (FDA Food Code %s)' % cod


MUTACIONES = [
    ('G1', F_VERDE, _bloquear_verde, 'desbloqueadas'),
    ('G1', F01, _quitar_tarea, 'tareas'),
    ('G1', F06, _area_rota, 'área de impresión'),
    ('G2', F01, lambda wb: _set(wb[mapas.HOJAS['Instrucciones']], 'B6', 'Revisa cada tarea con el equipo'), 'ES'),
    ('G2', F03, lambda wb: _set(wb[mapas.HOJAS['Instrucciones']], 'B6', 'Petty cash: 150 € per week'), 'moneda'),
    ('G2', F_HORAS, _hora_24, '24 h'),
    ('G2', F02, lambda wb: _set(wb[mapas.HOJAS['Instrucciones']], 'B6', 'Budget $5 per station'), '«$»'),
    ('G3', F03, lambda wb: setattr(wb[H_DIARIO].page_setup, 'paperSize', 9), 'paperSize'),
    ('G3', F01, lambda wb: setattr(wb[H_KO].oddFooter.center, 'text', 'Page &P'), 'pie'),
    ('G3', F01, _titulo_largo, 'ajuste'),
    ('G4', F01, _celda_glosario, 'mapa-glosario'),
    ('G4', F_HORAS, _hora_mal, 'hora_us'),
    ('G5', F09, _caja_mal, 'caja'),
    ('G6', F09, _cache_mal, 'número'),
    ('G7', F01, _grados_mal, '°F'),
    ('G7', F01, _cita_mal, 'cita normativa'),
]
if mapas.TERMINOS_PROHIBIDOS:                                    # vertical con términos prohibidos (cafe)
    MUTACIONES.append(('G2', F01, lambda wb: _set(wb[mapas.HOJAS['Instrucciones']], 'B6', (
        wb[mapas.HOJAS['Instrucciones']]['B6'].value or '') + ' ' + mapas.TERMINO_PROHIBIDO_EJEMPLO), 'término prohibido'))
if A.CF_POR_CELDA:                                               # vertical con CF cambiada a propósito (catering M5)
    _KCF, (_CF_ES, _CF_EN) = next(iter(A.CF_POR_CELDA.items()))

    def _cf_excepcion_mal(wb):
        """La excepción declarada vuelve a la fórmula ES: G1 tiene que cazarla (la excepción no es un pase libre)."""
        ws = wb[mapas.HOJAS[_KCF[1]]]
        cf = next(c for c in ws.conditional_formatting if str(c.sqref) == _KCF[2])
        cf.rules[0].formula = [_CF_ES]

    def _cf_ajena_mal(wb):
        """Una CF que NO es la excepción, alterada: G1 tiene que cazarla."""
        corto = next(f for f in mapas.LIBROS_ES if mapas.FICHEROS[f] == F09).split('-')[0]
        for ws in wb.worksheets:
            for cf in ws.conditional_formatting:
                h_es = next(h for h, e in mapas.HOJAS.items() if e == ws.title and h in mapas.HOJAS_ES_POR_LIBRO[corto])
                if (corto, h_es, str(cf.sqref)) not in A.CF_POR_CELDA and cf.rules and cf.rules[0].formula:
                    cf.rules[0].formula = [cf.rules[0].formula[0] + '*1']
                    return
        raise AssertionError('el 09 no tiene ninguna CF fuera de CF_POR_CELDA')

    MUTACIONES.append(('G1', mapas.FICHEROS[next(f for f in mapas.LIBROS_ES if f.startswith(_KCF[0] + '-'))],
                       _cf_excepcion_mal, 'formato condicional'))
    MUTACIONES.append(('G1', F09, _cf_ajena_mal, 'formato condicional'))
if mapas.VALORES_EN:                                             # vertical con números fijados (cafe: inventario)
    _KV, (_ES_V, _EN_V) = next(iter(mapas.VALORES_EN.items()))

    def _valor_mal(wb):
        wb[mapas.HOJAS[_KV[1]]][_KV[2]].value = _ES_V                # vuelve el texto ES («2»): G4 lo caza

    MUTACIONES.append(('G4', mapas.FICHEROS[next(f for f in mapas.LIBROS_ES if f.startswith(_KV[0]))], _valor_mal,
                       'valor fijado'))


def autotest(carpeta, mes, datos):
    ok = True
    base = ejecutar(carpeta, mes, ['G1', 'G2', 'G3', 'G4', 'G5', 'G6', 'G7'], datos, silencioso=True)
    sucios = [(g, v['fallos'][:3]) for g, v in base.gates.items() if v['fallos']]
    if sucios:
        print('AUTOTEST: la carpeta base no está limpia: arregla eso primero')
        for g, f in sucios:
            print('  %s: %s' % (g, f))
        return False
    for gate, fname, fn, aguja in MUTACIONES:
        tmp = tempfile.mkdtemp(prefix='checklists-autotest-')
        try:
            for f in mapas.FICHEROS.values():
                shutil.copy2(os.path.join(carpeta, f), tmp)
            _mutar(tmp, fname, fn)
            if gate == 'G6':
                recalcular(os.path.join(tmp, fname))
            R = ejecutar(tmp, mes, [gate], datos, silencioso=True)
            fallos = [f for g in R.gates.values() for f in g['fallos']]
            cazado = any(aguja.lower() in f.lower() for f in fallos)
            ok &= cazado
            print('  autotest %s %-48s %-18s %s' % (gate, fname, aguja, 'CAZADO' if cazado else 'NO CAZADO %r' % fallos[:2]))
        finally:
            shutil.rmtree(tmp)
    return ok


# ==========================================================================
# main
# ==========================================================================
GATES = OrderedDict([('G1', gateG1), ('G2', gateG2), ('G3', gateG3), ('G4', gateG4), ('G5', gateG5),
                     ('G6', gateG6), ('G7', gateG7), ('G8', gateG8)])
CON_MES = {'G3', 'G4'}


def cargar_datos(parcial=False, textos=A.TEXTOS_EN):
    tx, avisos = A.cargar_traducciones(textos, parcial)
    for a in avisos:
        print('  aviso: ' + a)
    return {'censo': json.load(open(os.path.join(kit.DATOS, 'censo_es.json'), encoding='utf-8')),
            'te': json.load(open(os.path.join(kit.DATOS, 'textos_es.json'), encoding='utf-8')),
            'tx': tx, 'parcial': parcial}


def ejecutar(carpeta, mes, solo, datos, silencioso=False):
    R = Resultado()
    for nombre, fn in GATES.items():
        if solo and nombre not in solo:
            continue
        A.termica()
        if nombre in CON_MES:
            fn(R, carpeta, datos, mes)
        else:
            fn(R, carpeta, datos)
        g = R.gates[R.actual]
        if not silencioso:
            estado = 'OK' if not g['fallos'] else 'FALLA (%d)' % len(g['fallos'])
            resumen = ' · '.join('%s=%s' % (k, v) for k, v in list(g['datos'].items())[:16] if not isinstance(v, list))
            print('%-24s %s  %s' % (R.actual, estado, resumen[:320]))
            for f in g['fallos'][:int(os.environ.get('CHECKLISTS_MAX_FALLOS', '15'))]:
                print('    - ' + f[:260])
    return R


def main():
    ap = argparse.ArgumentParser(description='Gates F2 del %s (EN)' % mapas.PRODUCTO)
    kit.argumento(ap)
    ap.add_argument('--dir', default=A.DESTINO_REAL, help='carpeta con los 11 xlsx EN (por defecto dl/<slug>)')
    ap.add_argument('--mes', default=None, help='mes de la línea de versión (por defecto, el actual)')
    ap.add_argument('--solo', default=None, help='G1,G2,… (por defecto, todos)')
    ap.add_argument('--json', default=None)
    ap.add_argument('--parcial', action='store_true', help='tolera la marca de relleno (dry-run --parcial)')
    ap.add_argument('--textos', default=A.TEXTOS_EN, help='carpeta de G<n>.json (la misma del montaje)')
    ap.add_argument('--autotest', action='store_true')
    args = ap.parse_args()
    mes = args.mes or A.mes_por_defecto()
    datos = cargar_datos(args.parcial, args.textos)
    carpeta = os.path.abspath(args.dir)
    if args.autotest:
        ok = autotest(carpeta, mes, datos)
        print('AUTOTEST: %s (%d defectos inyectados)' % ('OK' if ok else 'FALLA', len(MUTACIONES)))
        sys.exit(0 if ok else 1)
    solo = [s.strip().upper() for s in args.solo.split(',')] if args.solo else None
    R = ejecutar(carpeta, mes, solo, datos)
    fallos = sum(len(g['fallos']) for g in R.gates.values())
    if args.json:
        with open(args.json, 'w', encoding='utf-8') as fh:
            json.dump(R.gates, fh, ensure_ascii=False, indent=1, default=str)
    print('\n%s: %d gates, %d fallos (%s)%s' % ('VERDE' if not fallos else 'ROJO', len(R.gates), fallos, carpeta,
                                              ' · PARCIAL' if args.parcial else ''))
    sys.exit(1 if fallos else 0)


if __name__ == '__main__':
    main()
