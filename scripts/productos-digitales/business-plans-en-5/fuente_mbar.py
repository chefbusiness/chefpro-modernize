#!/usr/bin/env python3
"""
fuente_mbar.py — material de CONTENIDO para los redactores: el checklist (6 hojas), el plan de negocio, la carta de 15
cócteles, las 10 experiencias temáticas y el contrato de la coctelería para eventos ES (línea B,
`plan-negocio-cocteleria-eventos`) en `fuente_es_mbar.json`. El circuito de texto corre sobre el FOOD TRUCK ES (las
ranuras); el contenido de bar móvil sale de aquí y de research (D86, D91, D92). Sus CIFRAS no se usan (D93). Corre en
el VPS (openpyxl + python-docx). Sesión Claude Code, 13-oct-2026 (F2, paso 1).

    python3 bp5.py fuente
"""
import datetime
import hashlib
import json
import os
from collections import OrderedDict

import docx
import openpyxl

AQUI = os.path.dirname(os.path.abspath(__file__))
REPO = os.path.abspath(os.path.join(AQUI, '..', '..', '..'))
DIR = os.path.join(REPO, 'astro-site', 'public', 'dl', 'plan-negocio-cocteleria-eventos')
CHECK = 'checklist-apertura-cocteleria-eventos.xlsx'
PLAN = 'plan-de-negocio-cocteleria-eventos.docx'
CARTA = 'carta-15-cocktails-tematicos.docx'
EXPERIENCIAS = 'catalogo-10-experiencias-tematicas.docx'
CONTRATO = 'modelo-contrato-cocteleria-eventos.docx'
# F1-inventario-delta §5: tareas solo España (hoja, fila) — fuera del EN
SOLO_ES = {('F1 - Constitucion', 7), ('F1 - Constitucion', 8), ('F1 - Constitucion', 10), ('F1 - Constitucion', 13),
           ('F2 - Seguros', 8), ('F2 - Seguros', 11), ('F4 - Marca', 15)}
FASE_EN = OrderedDict([
    ('F1 - Constitucion', 'P1 Business Setup (+ P2 modelo legal / P4 manipulador)'),
    ('F2 - Seguros', 'P2 Trailer & Permits (seguros) (+ P4 OSHA / P5 privacy policy)'),
    ('F3 - Equipamiento', 'P3 Equipment (+ P2 remolque y DMV / P4 uniformes)'),
    ('F4 - Marca', 'P5 Marketing (+ P1 contrato y CRM)'),
    ('F5 - Captacion', 'P5 Marketing (+ P6 eventos piloto)'),
    ('F6 - Operativa', 'P6 First 90 Days (+ P1 cláusula de fotos / P4 formación)'),
])
USO = ('CONTENIDO de bar móvil para los redactores (D86, D91, D92). NO se traduce tal cual: se condensa y se adapta a '
       'EE. UU. con research (F1-research-us-delta.md §1, §5, §6, §8 y el glosario §9). Las CIFRAS del ES NO se usan '
       '(D93: 40.030, 70.000, 50 eventos/año, 1.400 €/evento, 42 %…): las del caso salen del libro EN (tokens). Fuera: '
       'tareas marcadas SOLO_ES, «España», «ChefBusiness», proveedores y directorios españoles (Bodas.net, Zankyou, '
       'CateringClick, GGM, Makro…), normativa española, importes en €, «freelance» (en US: W-2) y las erratas del ES. '
       'Recetas: ml → oz (30 ml = 1 oz) con paso de 0.25 oz; el escalado del ES está a la mitad (no lo copies)')


def sha(p):
    return hashlib.sha256(open(p, 'rb').read()).hexdigest()


def checklist():
    p = os.path.join(DIR, CHECK)
    wb = openpyxl.load_workbook(p, data_only=True)
    fases = OrderedDict()
    for ws in wb.worksheets:
        tareas = []
        titulo = ws['A2'].value
        for r in range(5, ws.max_row + 1):
            b, c, d, e = (ws.cell(r, k).value for k in range(2, 6))
            if not (isinstance(b, str) and b.strip()) or isinstance(c, (int, float)) or b.startswith('Tareas'):
                continue
            x = OrderedDict([('fila', r), ('tarea', b.strip()), ('responsable', c), ('plazo', d), ('nota', e)])
            if (ws.title, r) in SOLO_ES:
                x['marca'] = 'SOLO_ES (fuera)'
            tareas.append(x)
        fases[ws.title] = OrderedDict([('titulo', titulo), ('fase_en_sugerida', FASE_EN.get(ws.title, '')),
                                       ('tareas', tareas)])
    return p, fases


def parrafos(nombre, minimo=1):
    p = os.path.join(DIR, nombre)
    d = docx.Document(p)
    pars = [OrderedDict([('i', i), ('estilo', x.style.name), ('texto', x.text)]) for i, x in enumerate(d.paragraphs)
            if len(x.text.strip()) >= minimo and not set(x.text.strip()) <= set('━─—-· ')]
    tablas = [[[c.text for c in row.cells] for row in t.rows] for t in d.tables]
    return p, pars, tablas


def carta():
    """Las 15 recetas: nombre, temática, cristalería e ingredientes en ml (tabla k = receta k)."""
    p, pars, tablas = parrafos(CARTA)
    recetas, actual = [], None
    for k, x in enumerate(pars):
        t = x['texto'].strip()
        if k + 1 < len(pars) and pars[k + 1]['texto'].strip().startswith('Tematica:'):
            actual = OrderedDict([('nombre', t), ('ficha', []), ('ingredientes_ml', None)])
            recetas.append(actual)
        elif actual is not None and (t.startswith(('Tematica', 'Origen', 'Coste')) or len(t) > 60):
            if not t.startswith('Coste'):                                   # el coste en € no se usa (D93)
                actual['ficha'].append(t)
    for k, r in enumerate(recetas[:len(tablas)]):
        r['ingredientes_ml'] = [' | '.join(c for c in fila if c) for fila in tablas[k]]
    notas = [x['texto'] for x in pars if x['texto'].startswith('•')]
    return p, recetas, notas


def main():
    pc, fases = checklist()
    pp, pars, tablas = parrafos(PLAN)
    pk, recetas, notas_carta = carta()
    pe, exp, _t = parrafos(EXPERIENCIAS, minimo=3)
    pt, contrato, _t2 = parrafos(CONTRATO, minimo=3)
    out = OrderedDict([
        ('meta', OrderedDict([
            ('generado', datetime.datetime.now(datetime.timezone.utc).isoformat(timespec='seconds')),
            ('script', 'scripts/productos-digitales/business-plans-en-5/fuente_mbar.py'),
            ('uso', USO),
            ('fuentes', OrderedDict((os.path.relpath(x, REPO), sha(x)) for x in (pc, pp, pk, pe, pt))),
        ])),
        ('checklist_es', fases),
        ('plan_docx_es', OrderedDict([('parrafos', pars), ('tablas', tablas)])),
        ('carta_15_cocteles', OrderedDict([('recetas', recetas), ('notas', notas_carta)])),
        ('experiencias_10_tematicas', exp),
        ('contrato_es', contrato),
    ])
    dest = os.path.join(AQUI, 'fuente_es_mbar.json')
    with open(dest, 'w', encoding='utf-8') as fh:
        json.dump(out, fh, ensure_ascii=False, indent=1)
        fh.write('\n')
    n = sum(len(f['tareas']) for f in fases.values())
    fuera = sum(1 for f in fases.values() for t in f['tareas'] if 'marca' in t)
    print('fuente_es_mbar.json: %d fases, %d tareas (%d fuera) · plan %d párrafos, %d tablas · carta %d recetas · '
          'experiencias %d párrafos · contrato %d párrafos' % (len(fases), n, fuera, len(pars), len(tablas),
                                                              len(recetas), len(exp), len(contrato)))
    return 0


if __name__ == '__main__':
    raise SystemExit(main())
