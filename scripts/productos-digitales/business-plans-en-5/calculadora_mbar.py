#!/usr/bin/env python3
"""
calculadora_mbar.py — Mobile Bar Pricing Calculator (Cost per Guest to Quote + Shopping List), D89-D90 del Mobile Bar
Business Plan Kit (EN). Sobre el molde de `../business-plans-en-4/calculadora_cat.py` (mismas reglas: entradas verdes
con validación, salidas calculadas, el margen se CALCULA, ejemplos con las mismas fórmulas, docProps D20, US Letter,
sin anclas de precio, todo valor por defecto «Kit estimate» salvo cita). Cinco pestañas:

  1. Pricing Calculator      paquete (3: full bar – dry hire · limited bar – dry hire · full bar – licensed), invitados,
                             horas, copas (2 la 1.ª hora + 1/h: eventplanning.com), bartenders sugeridos por paquete,
                             costes → precio por invitado y por hora, subtotal, service charge, sales tax, total,
                             depósito, saldo y margen calculado.
  2. Drinks & Shopping List  copas por categoría según la mezcla del paquete, botellas de 750 ml (≈ 16 copas de 1.5 oz),
                             vino (5 copas por botella), cerveza, brindis, hielo en bolsas: la lista que el bar entrega
                             al anfitrión en dry hire, y lo que pone el bar.
  3. Signature Cocktail Specs 10 recetas de la carta ES de coctelería en oz (paso 0.25 oz), escaladas a 50/100/200 copas
                             POR FÓRMULA y con el coste por copa a partir del precio del envase que teclea el usuario.
  4. 3 Example Quotes        boda 120 inv. 5 h full bar · holiday party de empresa 100 inv. 4 h · cumpleaños privado 60
                             inv. 4 h limited bar: los 3 en dry hire y con las mismas fórmulas.
  5. Instructions

Corre en el VPS (openpyxl + pycel). Sesión Claude Code, 13-oct-2026 (F2, paso 1).

    python3 bp5.py calculadora              # genera + inject_cache + autotest
    python3 bp5.py calculadora --autotest   # solo el autotest sobre el fichero ya generado
"""
import argparse
import math
import os
import re
import sys
from collections import OrderedDict

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill, Protection
from openpyxl.worksheet.datavalidation import DataValidation

AQUI = os.path.dirname(os.path.abspath(__file__))
PD = os.path.dirname(AQUI)
REPO = os.path.abspath(os.path.join(AQUI, '..', '..', '..'))
sys.path.insert(0, os.path.join(PD, 'business-plans-en'))
sys.path.insert(0, PD)
os.environ.setdefault('BP_CONJUNTO', 'mbar')
import mapas as M                                                                  # noqa: E402

SLUG = 'mobile-bar-business-plan'
FICHERO = 'mobile-bar-pricing-calculator.xlsx'
TITULO = 'Mobile Bar Pricing Calculator (Cost per Guest to Quote + Shopping List)'
DESTINO = os.path.join(REPO, 'astro-site', 'public', 'dl', SLUG, FICHERO)
H1, H2, H3, H4, H5 = ('Pricing Calculator', 'Drinks & Shopping List', 'Signature Cocktail Specs', '3 Example Quotes',
                      'Instructions')
HOJAS_CALCULO = (H1, H2, H3, H4)
VERDE = PatternFill('solid', fgColor='E8F5E9')
GRIS = PatternFill('solid', fgColor='F2F2F2')
F = 'Calibri'
MONEDA, ENTERO, PCT, PCT1, DEC1 = '#,##0.00', '#,##0', '0%', '0.0%', '#,##0.0'
GLOSA = M.GLOSA_DRY_HIRE

FULL_DRY = 'Full bar – dry hire'
LIMITED_DRY = 'Limited bar (beer, wine + 1 signature) – dry hire'
FULL_LIC = 'Full bar – licensed (alcohol included)'
PAQUETES = [FULL_DRY, LIMITED_DRY, FULL_LIC]
# invitados por bartender según el paquete (theknot.com, resumen → Kit estimate): 1:50 barra completa, 1:75 limitada
RATIO = OrderedDict([(FULL_DRY, 50), (LIMITED_DRY, 75), (FULL_LIC, 50)])

# --------------------------------------------------------------------------
# 1. Pricing Calculator: entradas (clave, rótulo, formato, valor por defecto, nota). Fila = 5 + índice
# --------------------------------------------------------------------------
ENTRADAS = [
    ('package', 'Bar package', '@', FULL_DRY,
     'Choose from the list. The licensed package is ONLY for a business that holds a liquor license or caterer\'s '
     'permit: see section 9 of the business plan. In dry hire the host buys and owns the alcohol.'),
    ('guests', 'Guests expected', ENTERO, 95, 'Whole number of adult guests at the event.'),
    ('hours', 'Hours of bar service', DEC1, 4, 'Service time only; setup, travel and breakdown go two rows below.'),
    ('first_hour', 'Drinks per guest in the first hour', DEC1, 2,
     'Planning rule: 2 drinks per guest in the first hour (eventplanning.com).'),
    ('per_hour', 'Drinks per guest in each following hour', DEC1, 1,
     'Planning rule: 1 drink per guest per hour after the first (eventplanning.com).'),
    ('bartenders', 'Bartenders working the event', ENTERO, 2,
     'Compare with the suggested number in the results (table at the bottom).'),
    ('barbacks', 'Barbacks / setup crew working the event', ENTERO, 0, 'Usually needed at large weddings.'),
    ('extra_hours', 'Setup, travel and breakdown hours per person', DEC1, 3,
     'Paid hours on top of service: loading, travel, setup, breakdown. Kit estimate.'),
    ('wage_bt', 'Bartender hourly wage', MONEDA, 20.0,
     'Kit estimate. BLS (May 2025): mean 19.61 an hour, median 16.51, tips included. Bartenders are W-2 employees.'),
    ('wage_bb', 'Barback hourly wage', MONEDA, 16.0, 'Kit estimate. At or above your state minimum wage.'),
    ('payroll', "Employer payroll taxes and workers' comp (% of wages)", PCT, 0.10, 'Kit estimate.'),
    ('mixers', 'Mixers, garnishes and N/A drinks per drink served', MONEDA, 0.38,
     'Kit estimate. Juices, syrups, sodas, fruit and mocktail ingredients: what the bar supplies.'),
    ('alcohol', 'Alcohol cost per drink (licensed package only)', MONEDA, 3.0,
     'Used ONLY with the licensed package. In dry hire the host buys the alcohol: it is never in your cost or your '
     'price.'),
    ('cups', 'Cups, straws, napkins and menu signs per guest', MONEDA, 0.55, 'Kit estimate.'),
    ('ice_lb', 'Ice per guest (lb)', DEC1, 1.5, 'Kit estimate: 1.5 lb per guest is a common planning rule.'),
    ('ice_price', 'Bagged ice cost per lb', MONEDA, 0.25, 'Kit estimate: use your supplier\'s price.'),
    ('miles', 'Round-trip miles', ENTERO, 40, 'Storage unit → venue → storage unit.'),
    ('per_mile', 'Cost per mile (tow vehicle + trailer)', MONEDA, 0.70,
     'Kit estimate: use your real cost per mile or the current IRS standard mileage rate.'),
    ('bar_fee', 'Flat bar fee charged per event (trailer or bar setup)', MONEDA, 0.0,
     'Optional flat charge on top of the per-guest price ("a flat fee plus a price per guest"). 0 = none.'),
    ('upcharge', 'Specialty bar upcharge (% of the per-guest price)', PCT, 0.0,
     'Mimosa bar, espresso martini bar and other themed bars. 0 = none.'),
    ('commission', 'Marketplace or planner fee on this booking (% of subtotal)', PCT, 0.0,
     'GigSalad charges 5% on the free plan or 2.5% on a paid plan per confirmed booking (gigsalad.com). 0 = direct '
     'booking.'),
    ('margin', 'Target gross margin on the bar service subtotal', PCT, 0.55,
     'Share of the subtotal left after the direct costs above and the marketplace fee. Kit estimate.'),
    ('sc', 'Service charge (% of subtotal)', PCT, 0.0,
     'Added on top. The part you pay out to staff is wages, not tips (IRS Rev. Rul. 2012-18), with payroll taxes.'),
    ('tax', 'Sales tax rate (state + local)', PCT1, 0.08, 'Kit estimate: use your own state + local rate.'),
    ('tax_sc', 'Does your state tax the bar service and the service charge?', '@', 'Yes',
     'California taxes a caterer\'s service even when the client supplies the drinks (cdtfa.ca.gov). Rules vary by '
     'state: check yours.'),
    ('deposit', 'Deposit to book the date (% of total)', PCT, 0.30,
     'A 30% deposit is typical (checkcherry.com); the balance is due before the event.'),
    ('minimum', 'Minimum charge for the bar service (subtotal floor)', MONEDA, 400.0,
     'Below it you bill the minimum. Kit estimate. 0 = no minimum.'),
    ('extra_hour_price', 'Price of an extra hour of service (quoted separately)', MONEDA, 150.0,
     'Kit estimate. It is listed in the quote, not added to the total.'),
]
K = {k: 5 + i for i, (k, *_r) in enumerate(ENTRADAS)}                  # clave → fila de la hoja 1
F_CALC = 5 + len(ENTRADAS) + 1                                         # cabecera del bloque de cálculo
SALIDAS = [
    ('drinks', 'Drinks to plan for', ENTERO, '{guests}*({first_hour}+max(0,{hours}-1)*{per_hour})',
     'Guests × (first-hour drinks + following hours × drinks per hour).'),
    ('ice_total', 'Ice to bring (lb)', DEC1, '{guests}*{ice_lb}', ''),
    ('bt_sug', 'Suggested bartenders for this package (Kit estimate)', ENTERO, 'sugerido',
     'Guests ÷ guests per bartender in the table below, rounded up.'),
    ('c_staff', 'Staff cost (wages + payroll taxes)', MONEDA,
     '({bartenders}*{wage_bt}+{barbacks}*{wage_bb})*({hours}+{extra_hours})*(1+{payroll})', ''),
    ('c_mixers', 'Mixers, garnishes and N/A drinks', MONEDA, '{drinks}*{mixers}', ''),
    ('c_alcohol', 'Alcohol (licensed package only)', MONEDA, 'licencia', 'Zero in dry hire: the host buys the alcohol.'),
    ('c_cups', 'Cups, straws, napkins and menu signs', MONEDA, '{guests}*{cups}', ''),
    ('c_ice', 'Ice', MONEDA, '{ice_total}*{ice_price}', ''),
    ('c_travel', 'Mileage', MONEDA, '{miles}*{per_mile}', ''),
    ('c_total', 'Total direct cost', MONEDA, '{c_staff}+{c_mixers}+{c_alcohol}+{c_cups}+{c_ice}+{c_travel}', ''),
    ('c_guest', 'Direct cost per guest', MONEDA, '{c_total}/{guests}', ''),
    ('price', 'PRICE PER GUEST before service charge and tax', MONEDA,
     'roundup(max(0,{c_total}/(1-{margin}-{commission})-{bar_fee})/{guests}*(1+{upcharge})*2)/2',
     'Rounded up to the next 0.50. Staff, mixers, ice and cups are inside it; the host\'s alcohol is not.'),
    ('subtotal', 'Bar service subtotal', MONEDA, 'max({price}*{guests}+{bar_fee},{minimum})',
     'Price per guest × guests + flat bar fee, never below the minimum.'),
    ('per_hour_eq', 'Equivalent price per hour of service', MONEDA, '{subtotal}/{hours}',
     'The same quote expressed per hour, for clients who compare hourly rates.'),
    ('sc_amt', 'Service charge', MONEDA, '{subtotal}*{sc}', ''),
    ('tax_amt', 'Sales tax', MONEDA, '({subtotal}+if({tax_sc}=="Yes",{sc_amt},0))*{tax}',
     'On the subtotal, plus the service charge if your state taxes it.'),
    ('total', 'TOTAL TO QUOTE', MONEDA, '{subtotal}+{sc_amt}+{tax_amt}', ''),
    ('dep_amt', 'Deposit to book the date', MONEDA, '{total}*{deposit}', ''),
    ('balance', 'Balance due before the event', MONEDA, '{total}-{dep_amt}', ''),
    ('comm_amt', 'Marketplace or planner fee', MONEDA, '{subtotal}*{commission}', 'What the marketplace keeps.'),
    ('profit', 'Gross profit on the subtotal', MONEDA, '{subtotal}-{comm_amt}-{c_total}', ''),
    ('margin_calc', 'Gross margin (calculated)', PCT1, 'if({subtotal}==0,0,{profit}/{subtotal})',
     'Calculated from the costs above, never typed. It is at or above the target because the price is rounded up.'),
    ('extra_hour', 'Extra hour of service (not in the total)', MONEDA, '{extra_hour_price}', ''),
]
S = {k: F_CALC + 1 + i for i, (k, *_r) in enumerate(SALIDAS)}
F_TABLA = F_CALC + 1 + len(SALIDAS) + 1                                # cabecera de la tabla de paquetes
T_NOMBRES = "'%s'!$A$%d:$A$%d" % (H1, F_TABLA + 1, F_TABLA + len(RATIO))
T_RATIOS = "'%s'!$B$%d:$B$%d" % (H1, F_TABLA + 1, F_TABLA + len(RATIO))
T_LICENCIA = "'%s'!$A$%d" % (H1, F_TABLA + 3)

EJEMPLOS = OrderedDict([
    ('Wedding, 120 guests, 5 hours', dict(
        package=FULL_DRY, guests=120, hours=5, first_hour=2, per_hour=1, bartenders=3, barbacks=1, extra_hours=3.5,
        wage_bt=20.0, wage_bb=16.0, payroll=0.10, mixers=0.38, alcohol=3.0, cups=0.55, ice_lb=1.5, ice_price=0.25,
        miles=60, per_mile=0.70, bar_fee=250.0, upcharge=0.0, commission=0.0, margin=0.55, sc=0.15, tax=0.08,
        tax_sc='Yes', deposit=0.30, minimum=400.0, extra_hour_price=150.0)),
    ('Corporate holiday party, 100 guests, 4 hours', dict(
        package=FULL_DRY, guests=100, hours=4, first_hour=2, per_hour=1, bartenders=2, barbacks=1, extra_hours=3,
        wage_bt=20.0, wage_bb=16.0, payroll=0.10, mixers=0.38, alcohol=3.0, cups=0.55, ice_lb=1.5, ice_price=0.25,
        miles=30, per_mile=0.70, bar_fee=0.0, upcharge=0.0, commission=0.0, margin=0.55, sc=0.0, tax=0.08,
        tax_sc='Yes', deposit=0.30, minimum=400.0, extra_hour_price=150.0)),
    ('Private birthday, 60 guests, 4 hours', dict(
        package=LIMITED_DRY, guests=60, hours=4, first_hour=2, per_hour=1, bartenders=1, barbacks=0, extra_hours=2.5,
        wage_bt=20.0, wage_bb=16.0, payroll=0.10, mixers=0.38, alcohol=3.0, cups=0.55, ice_lb=1.5, ice_price=0.25,
        miles=20, per_mile=0.70, bar_fee=0.0, upcharge=0.15, commission=0.05, margin=0.55, sc=0.0, tax=0.08,
        tax_sc='No', deposit=0.30, minimum=400.0, extra_hour_price=150.0)),
])
EJ_NOTAS = OrderedDict([
    ('Wedding, 120 guests, 5 hours', 'Dry hire full bar from the trailer: 3 bartenders and a barback, a flat bar fee for '
                                     'the trailer and a 15% service charge; 30% deposit to book the date.'),
    ('Corporate holiday party, 100 guests, 4 hours', 'Dry hire full bar on a weeknight in December, booked directly by '
                                                     'the company: no service charge, no marketplace fee.'),
    ('Private birthday, 60 guests, 4 hours', 'Dry hire limited bar (beer, wine and one signature cocktail) with a '
                                             'mimosa bar upcharge, booked through a marketplace (5% fee); this '
                                             'example assumes a state that does not tax the service (switch set to '
                                             'No).'),
])


# --------------------------------------------------------------------------
# 2. Drinks & Shopping List: parámetros (clave, rótulo, formato, valor, nota) y mezcla por paquete
# --------------------------------------------------------------------------
MEZCLA = OrderedDict([(FULL_DRY, (0.50, 0.25, 0.25)), (LIMITED_DRY, (0.20, 0.40, 0.40)), (FULL_LIC, (0.50, 0.25, 0.25))])
COMPRA = [
    ('pours', 'Drinks per 750 ml bottle of spirits (1.5 oz pours)', ENTERO, 16,
     'A 750 ml bottle holds about 25 oz: about 16 pours of 1.5 oz.'),
    ('wine_glasses', 'Glasses per 750 ml bottle of wine', ENTERO, 5, 'About 5 glasses of 5 oz.'),
    ('beer_per', 'Beers (12 oz cans or bottles) per beer drink', DEC1, 1, ''),
    ('buffer', 'Safety buffer on every quantity', PCT, 0.10, 'Kit estimate: unopened bottles stay with the host.'),
    ('toast', 'Champagne toast?', '@', 'No', 'Yes or No.'),
    ('toast_pours', 'Toast pours per 750 ml bottle of sparkling wine', ENTERO, 8, 'Kit estimate.'),
    ('bag_lb', 'Bag of ice (lb)', ENTERO, 10, 'Use the bag size your supplier sells.'),
    ('cups_per', 'Cups per drink', DEC1, 1, ''),
    ('napkins_per', 'Napkins per drink', DEC1, 1, ''),
]
LICORES = OrderedDict([('Vodka', 0.30), ('Tequila', 0.20), ('Whiskey / bourbon', 0.20), ('Rum', 0.15), ('Gin', 0.15)])


# --------------------------------------------------------------------------
# 3. Signature Cocktail Specs: 10 recetas de la carta ES (ml → oz, paso 0.25 oz). Ingrediente: (nombre, cantidad por
#    copa, unidad, quién lo compra en dry hire, tamaño del envase, unidad del envase, precio del envase [Kit estimate])
#    Todo lo que lleva alcohol (bitters incluidos) es del anfitrión en dry hire.
# --------------------------------------------------------------------------
_B, _H = 'Bar', 'Host'
_L750, _L1 = 25.4, 33.8                                                # 750 ml y 1 L en oz
RECETAS = [
    ('Old Fashioned', 'Rocks glass, one large cube', '', [
        ('Bourbon', 2, 'oz', _H, _L750, 'oz', 25.0), ('Angostura bitters', 2, 'dash', _H, 120, 'dash', 10.0),
        ('Sugar cube', 1, 'each', _B, 126, 'each', 4.0), ('Orange peel', 1, 'each', _B, 4, 'each', 1.0)]),
    ('Negroni', 'Rocks glass, one large cube', '', [
        ('Gin', 1, 'oz', _H, _L750, 'oz', 25.0), ('Sweet vermouth', 1, 'oz', _H, _L750, 'oz', 18.0),
        ('Campari', 1, 'oz', _H, _L750, 'oz', 28.0), ('Orange peel', 1, 'each', _B, 4, 'each', 1.0)]),
    ('French 75', 'Flute or coupe', '', [
        ('Gin', 1, 'oz', _H, _L750, 'oz', 25.0), ('Fresh lemon juice', 0.5, 'oz', _B, 32, 'oz', 10.0),
        ('Simple syrup', 0.5, 'oz', _B, _L1, 'oz', 8.0), ('Sparkling wine (brut)', 2, 'oz', _H, _L750, 'oz', 14.0),
        ('Lemon twist', 1, 'each', _B, 6, 'each', 0.75)]),
    ('Mai Tai', 'Rocks glass or tiki mug, crushed ice', 'Allergen: orgeat is made from almonds (tree nuts)', [
        ('Aged dark rum', 1, 'oz', _H, _L750, 'oz', 25.0), ('White rum', 1, 'oz', _H, _L750, 'oz', 20.0),
        ('Orange curacao', 0.5, 'oz', _H, _L750, 'oz', 22.0), ('Orgeat (almond) syrup', 0.5, 'oz', _B, _L750, 'oz', 12.0),
        ('Fresh lime juice', 0.75, 'oz', _B, 32, 'oz', 10.0), ('Mint sprig', 1, 'each', _B, 20, 'each', 2.0)]),
    ('Pina Colada', 'Hurricane glass', 'Allergen: coconut (cream of coconut)', [
        ('White rum', 2, 'oz', _H, _L750, 'oz', 20.0), ('Aged rum float', 0.5, 'oz', _H, _L750, 'oz', 25.0),
        ('Cream of coconut', 2, 'oz', _B, 15, 'oz', 4.0), ('Pineapple juice', 4, 'oz', _B, 46, 'oz', 4.0),
        ('Pineapple wedge', 1, 'each', _B, 16, 'each', 4.0)]),
    ('Passion Fruit Mojito', 'Collins glass, crushed ice', '', [
        ('White rum', 1.75, 'oz', _H, _L750, 'oz', 20.0), ('Passion fruit puree', 1, 'oz', _B, 32, 'oz', 15.0),
        ('Fresh lime juice', 0.75, 'oz', _B, 32, 'oz', 10.0), ('Simple syrup', 0.5, 'oz', _B, _L1, 'oz', 8.0),
        ('Mint leaves', 8, 'each', _B, 60, 'each', 2.0), ('Club soda (top)', 2, 'oz', _B, _L1, 'oz', 1.5)]),
    ('Aperol Spritz', 'Wine glass', '', [
        ('Aperol', 2, 'oz', _H, _L750, 'oz', 25.0), ('Prosecco', 3, 'oz', _H, _L750, 'oz', 14.0),
        ('Club soda', 1, 'oz', _B, _L1, 'oz', 1.5), ('Orange slice', 1, 'each', _B, 8, 'each', 1.0)]),
    ('Espresso Martini', 'Chilled martini glass or coupe', '', [
        ('Vodka', 1.75, 'oz', _H, _L750, 'oz', 25.0), ('Coffee liqueur', 0.75, 'oz', _H, _L750, 'oz', 25.0),
        ('Fresh espresso (cooled)', 1, 'oz', _B, 1, 'oz', 0.5), ('Simple syrup', 0.25, 'oz', _B, _L1, 'oz', 8.0),
        ('Coffee beans (garnish)', 3, 'each', _B, 300, 'each', 5.0)]),
    ('Smoked Mezcal Margarita', 'Coupe with a smoked-salt rim', '', [
        ('Mezcal', 1.75, 'oz', _H, _L750, 'oz', 40.0), ('Triple sec', 0.75, 'oz', _H, _L750, 'oz', 20.0),
        ('Fresh lime juice', 0.75, 'oz', _B, 32, 'oz', 10.0), ('Simple syrup', 0.25, 'oz', _B, _L1, 'oz', 8.0),
        ('Smoked salt (rim)', 1, 'each', _B, 60, 'each', 6.0), ('Lime wedge', 1, 'each', _B, 8, 'each', 0.5)]),
    ('Tropical Mocktail (zero-proof)', 'Hurricane glass, crushed ice', 'Allergen: coconut (cream of coconut)', [
        ('Pineapple juice', 3, 'oz', _B, 46, 'oz', 4.0), ('Fresh orange juice', 2, 'oz', _B, 52, 'oz', 5.0),
        ('Fresh lime juice', 0.75, 'oz', _B, 32, 'oz', 10.0), ('Grenadine', 0.5, 'oz', _B, _L750, 'oz', 8.0),
        ('Cream of coconut', 0.75, 'oz', _B, 15, 'oz', 4.0), ('Club soda (top)', 2, 'oz', _B, _L1, 'oz', 1.5)]),
]
ESCALAS = (50, 100, 200)


# --------------------------------------------------------------------------
# Lógica en Python (la misma que las fórmulas): la usa el autotest como valor esperado
# --------------------------------------------------------------------------
def calcular(e):
    r = OrderedDict()
    r['drinks'] = e['guests'] * (e['first_hour'] + max(0, e['hours'] - 1) * e['per_hour'])
    r['ice_total'] = e['guests'] * e['ice_lb']
    r['bt_sug'] = math.ceil(round(e['guests'] / RATIO[e['package']], 9))
    r['c_staff'] = ((e['bartenders'] * e['wage_bt'] + e['barbacks'] * e['wage_bb']) * (e['hours'] + e['extra_hours'])
                    * (1 + e['payroll']))
    r['c_mixers'] = r['drinks'] * e['mixers']
    r['c_alcohol'] = r['drinks'] * e['alcohol'] if e['package'] == FULL_LIC else 0
    r['c_cups'] = e['guests'] * e['cups']
    r['c_ice'] = r['ice_total'] * e['ice_price']
    r['c_travel'] = e['miles'] * e['per_mile']
    r['c_total'] = r['c_staff'] + r['c_mixers'] + r['c_alcohol'] + r['c_cups'] + r['c_ice'] + r['c_travel']
    r['c_guest'] = r['c_total'] / e['guests']
    base = max(0, r['c_total'] / (1 - e['margin'] - e['commission']) - e['bar_fee'])
    r['price'] = math.ceil(round(base / e['guests'] * (1 + e['upcharge']) * 2, 9)) / 2
    r['subtotal'] = max(r['price'] * e['guests'] + e['bar_fee'], e['minimum'])
    r['per_hour_eq'] = r['subtotal'] / e['hours']
    r['sc_amt'] = r['subtotal'] * e['sc']
    r['tax_amt'] = (r['subtotal'] + (r['sc_amt'] if e['tax_sc'] == 'Yes' else 0)) * e['tax']
    r['total'] = r['subtotal'] + r['sc_amt'] + r['tax_amt']
    r['dep_amt'] = r['total'] * e['deposit']
    r['balance'] = r['total'] - r['dep_amt']
    r['comm_amt'] = r['subtotal'] * e['commission']
    r['profit'] = r['subtotal'] - r['comm_amt'] - r['c_total']
    r['margin_calc'] = 0 if r['subtotal'] == 0 else r['profit'] / r['subtotal']
    r['extra_hour'] = e['extra_hour_price']
    return r


def calcular_compra(e, c=None, licores=LICORES, mezcla=MEZCLA):
    """Drinks & Shopping List con las entradas de la hoja 1 (e) y los parámetros de compra (c)."""
    c = c or OrderedDict((k, v) for k, _r, _f, v, _n in COMPRA)
    r = OrderedDict()
    drinks = calcular(e)['drinks']
    sc, sw, sb = mezcla[e['package']]
    up = 1 + c['buffer']
    r['d_cocktails'] = drinks * sc
    r['d_wine'] = drinks * sw
    r['d_beer'] = drinks * sb
    r['b_spirits'] = math.ceil(round(r['d_cocktails'] * up / c['pours'], 9))
    for i, (n, s) in enumerate(licores.items()):
        r['b_lic_%d' % i] = math.ceil(round(r['d_cocktails'] * s * up / c['pours'], 9))
    r['b_wine'] = math.ceil(round(r['d_wine'] * up / c['wine_glasses'], 9))
    r['u_beer'] = math.ceil(round(r['d_beer'] * c['beer_per'] * up, 9))
    r['b_toast'] = math.ceil(round(e['guests'] / c['toast_pours'], 9)) if c['toast'] == 'Yes' else 0
    r['ice_lb'] = e['guests'] * e['ice_lb']
    r['ice_bags'] = math.ceil(round(r['ice_lb'] / c['bag_lb'], 9))
    r['cups_n'] = math.ceil(round(drinks * c['cups_per'] * up, 9))
    r['napkins_n'] = math.ceil(round(drinks * c['napkins_per'] * up, 9))
    r['mixer_servings'] = math.ceil(round(r['d_cocktails'] * up, 9))
    return r


def calcular_specs():
    """Por receta: coste por copa total y de lo que pone el bar; por ingrediente, coste, escalado y envases."""
    out = []
    for nombre, _vaso, _al, ings in RECETAS:
        filas = []
        for ing, cant, _u, quien, env, _ue, precio in ings:
            filas.append(OrderedDict([('coste', cant / env * precio), ('x50', cant * 50), ('x100', cant * 100),
                                      ('x200', cant * 200), ('envases100', math.ceil(round(cant * 100 / env, 9))),
                                      ('quien', quien)]))
        tot = sum(f['coste'] for f in filas)
        bar = sum(f['coste'] for f in filas if f['quien'] == _B)
        out.append((nombre, filas, tot, bar))
    return out


def a_excel(expr, ref):
    """Plantilla Python-ish → fórmula Excel con las celdas de `ref` (clave → coordenada)."""
    f = expr.replace('max(', 'MAX(').replace('roundup(', 'ROUNDUP(').replace('if(', 'IF(').replace('==', '=')
    f = re.sub(r'\{(\w+)\}', lambda m: ref[m.group(1)], f)
    return f.replace('*2)/2', '*2,0)/2')


def formula_sugerido(ref_paquete, ref_guests):
    rat = 'INDEX(%s,MATCH(%s,%s,0))' % (T_RATIOS, ref_paquete, T_NOMBRES)
    return '=IFERROR(ROUNDUP(%s/%s,0),"—")' % (ref_guests, rat)


def formula_salida(k, expr, ref):
    if expr == 'sugerido':
        return formula_sugerido(ref['package'], ref['guests'])
    if expr == 'licencia':
        return '=IF(%s=%s,%s*%s,0)' % (ref['package'], T_LICENCIA, ref['drinks'], ref['alcohol'])
    return '=' + a_excel(expr, ref)


# --------------------------------------------------------------------------
# Construcción
# --------------------------------------------------------------------------
def _f(ws, coord, v, b=False, sz=11, i=False, fmt=None, wrap=False, fill=None, locked=True, color=None):
    c = ws[coord]
    c.value = v
    c.font = Font(name=F, size=sz, bold=b, italic=i, color=color)
    if fmt:
        c.number_format = fmt
    c.alignment = Alignment(wrap_text=wrap, vertical='top')
    if fill:
        c.fill = fill
    c.protection = Protection(locked=locked)
    return c


def _imprimir(ws, horizontal=False):
    ws.page_setup.paperSize = M.PAPEL_EN                                         # 1 = US Letter
    ws.page_setup.orientation = 'landscape' if horizontal else 'portrait'
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.oddFooter.center.text = M.FOOTER_EN
    ws.print_options.horizontalCentered = True
    ws.protection.sheet = True                                                   # sin contraseña (D20)
    ws.sheet_view.showGridLines = False


def _dv(ws, tipo, celdas, f1=None, f2=None, op=None, titulo='Invalid value', msg=''):
    dv = DataValidation(type=tipo, formula1=f1, formula2=f2, operator=op, allow_blank=False, showErrorMessage=True,
                        errorTitle=titulo, error=msg)
    ws.add_data_validation(dv)
    for c in celdas:
        dv.add(c)


def _cabecera(ws, titulo, intro, ancho_merge):
    _f(ws, 'A1', titulo, b=True, sz=14)
    _f(ws, 'A2', intro, wrap=True)
    ws.merge_cells('A2:%s2' % ancho_merge)
    ws.row_dimensions[2].height = 48


def hoja_calculadora(wb):
    ws = wb.active
    ws.title = H1
    for col, w in (('A', 60), ('B', 22), ('C', 80)):
        ws.column_dimensions[col].width = w
    _cabecera(ws, 'MOBILE BAR PRICING CALCULATOR — COST PER GUEST TO QUOTE',
              'Type in the GREEN cells; everything else is calculated. The default model is dry hire ' + GLOSA + ': '
              'the host buys and owns the alcohol, so it is never in your cost or your price. Amounts are in your '
              'currency and exclude sales tax unless the row says otherwise. Every default value is a kit estimate: '
              'replace it with your own numbers.', 'C')
    for col, t in zip('ABC', ('INPUTS', 'Value', 'Notes')):
        _f(ws, '%s4' % col, t, b=True, fill=GRIS)
    for k, rot, fmt, val, nota in ENTRADAS:
        f = K[k]
        _f(ws, 'A%d' % f, rot)
        _f(ws, 'B%d' % f, val, fmt=fmt if fmt != '@' else None, fill=VERDE, locked=False, wrap=k == 'package')
        _f(ws, 'C%d' % f, nota, wrap=True)
        if len(nota) > 90 or k == 'package':
            ws.row_dimensions[f].height = 15 * math.ceil(len(nota) / 90.0) + 2
    for col, t in zip('ABC', ('CALCULATION', 'Value', 'Notes')):
        _f(ws, '%s%d' % (col, F_CALC), t, b=True, fill=GRIS)
    ref = {k: '$B$%d' % f for k, f in K.items()}
    ref.update({k: '$B$%d' % f for k, f in S.items()})
    for k, rot, fmt, expr, nota in SALIDAS:
        f = S[k]
        neg = k in ('price', 'total', 'margin_calc')
        _f(ws, 'A%d' % f, rot, b=neg)
        _f(ws, 'B%d' % f, formula_salida(k, expr, ref), fmt=fmt, b=neg)
        if nota:
            _f(ws, 'C%d' % f, nota, wrap=True)
    for col, t in zip('ABC', ('GUESTS PER BARTENDER BY PACKAGE (Kit estimate — edit the numbers)', 'Guests', 'Notes')):
        _f(ws, '%s%d' % (col, F_TABLA), t, b=True, fill=GRIS)
    notas = {FULL_DRY: 'About 1 bartender per 50 guests with a full bar.',
             LIMITED_DRY: 'About 1 bartender per 75 guests with beer, wine and one signature cocktail.',
             FULL_LIC: 'Only if you hold the license (plan section 9): you also buy and serve the alcohol.'}
    for i, (paq, n) in enumerate(RATIO.items(), 1):
        _f(ws, 'A%d' % (F_TABLA + i), paq)
        _f(ws, 'B%d' % (F_TABLA + i), n, fmt=ENTERO, fill=VERDE, locked=False)
        _f(ws, 'C%d' % (F_TABLA + i), notas[paq], wrap=True)
    # validaciones (D89)
    _dv(ws, 'list', [ref['package'].replace('$', '')], f1=T_NOMBRES.replace("'%s'!" % H1, ''),
        titulo='Invalid package', msg='Choose a bar package from the list.')
    _dv(ws, 'list', [ref['tax_sc'].replace('$', '')], f1='"Yes,No"', titulo='Invalid entry', msg='Choose Yes or No.')
    _dv(ws, 'whole', ['B%d' % K['guests']], f1='1', f2='100000', op='between', titulo='Invalid count',
        msg='Enter a whole number of 1 or more.')
    _dv(ws, 'whole', ['B%d' % K[k] for k in ('bartenders', 'barbacks', 'miles')], f1='0', f2='100000', op='between',
        titulo='Invalid count', msg='Enter a whole number of 0 or more.')
    _dv(ws, 'whole', ['B%d' % (F_TABLA + i) for i in range(1, len(RATIO) + 1)], f1='1', f2='1000', op='between',
        titulo='Invalid count', msg='Enter a whole number of guests per bartender, 1 or more.')
    _dv(ws, 'decimal', ['B%d' % K['hours']], f1='0.5', f2='24', op='between', titulo='Invalid hours',
        msg='Enter the hours of service, between 0.5 and 24.')
    _dv(ws, 'decimal', ['B%d' % K[k] for k in ('first_hour', 'per_hour', 'extra_hours', 'wage_bt', 'wage_bb', 'mixers',
                                              'alcohol', 'cups', 'ice_lb', 'ice_price', 'per_mile', 'bar_fee',
                                              'minimum', 'extra_hour_price')],
        f1='0', op='greaterThanOrEqual', titulo='Invalid amount', msg='Enter a number greater than or equal to 0.')
    _dv(ws, 'decimal', ['B%d' % K[k] for k in ('payroll', 'upcharge', 'sc', 'tax', 'deposit')], f1='0', f2='1',
        op='between', titulo='Invalid percentage', msg='Enter a percentage between 0 and 1 (0.20 = 20%).')
    _dv(ws, 'decimal', ['B%d' % K['commission']], f1='0', f2='0.3', op='between', titulo='Invalid percentage',
        msg='Enter a fee between 0 and 30% (0.05 = 5%).')
    _dv(ws, 'decimal', ['B%d' % K['margin']], f1='0', f2='0.65', op='between', titulo='Invalid percentage',
        msg='Enter a target margin between 0 and 65% (0.55 = 55%).')
    ws.freeze_panes = 'A5'
    _imprimir(ws)
    return ws


def hoja_compra(wb):
    ws = wb.create_sheet(H2)
    for col, w in (('A', 56), ('B', 18), ('C', 18), ('D', 18), ('E', 60)):
        ws.column_dimensions[col].width = w
    _cabecera(ws, 'DRINKS & SHOPPING LIST — WHAT THE HOST BUYS AND WHAT THE BAR BRINGS',
              'It reads the package, guests, hours and drinks from the "%s" tab. In dry hire this is the shopping list '
              'you hand the host about two weeks before the event: the host buys and owns the alcohol, and unopened '
              'bottles stay with the host. With the licensed package, you buy it. Green cells are kit estimates: edit '
              'them.' % H1, 'E')
    g = "'%s'!" % H1
    fila = 4
    _f(ws, 'A%d' % fila, 'THE EVENT (from the "%s" tab)' % H1, b=True, fill=GRIS)
    ev = OrderedDict([('package', ('Bar package', g + '$B$%d' % K['package'], None)),
                      ('guests', ('Guests', g + '$B$%d' % K['guests'], ENTERO)),
                      ('hours', ('Hours of bar service', g + '$B$%d' % K['hours'], DEC1)),
                      ('drinks', ('Drinks to plan for', g + '$B$%d' % S['drinks'], ENTERO))])
    R = {}
    for k, (rot, src, fmt) in ev.items():
        fila += 1
        _f(ws, 'A%d' % fila, rot)
        _f(ws, 'B%d' % fila, '=' + src, fmt=fmt)
        R[k] = '$B$%d' % fila
    ws.merge_cells('B%d:D%d' % (int(R['package'][3:]), int(R['package'][3:])))
    fila += 2
    _f(ws, 'A%d' % fila, 'DRINK MIX BY PACKAGE (share of drinks; Kit estimate — edit)', b=True, fill=GRIS)
    for col, t in zip('BCD', ('Cocktails & spirits', 'Wine', 'Beer')):
        _f(ws, '%s%d' % (col, fila), t, b=True, fill=GRIS, wrap=True)
    _f(ws, 'E%d' % fila, 'Check (must be 100%)', b=True, fill=GRIS)
    f_mix0 = fila + 1
    verdes, pct_celdas = [], []
    for paq, (a, b, c) in MEZCLA.items():
        fila += 1
        _f(ws, 'A%d' % fila, paq, wrap=True)
        for col, v in zip('BCD', (a, b, c)):
            _f(ws, '%s%d' % (col, fila), v, fmt=PCT, fill=VERDE, locked=False)
            pct_celdas.append('%s%d' % (col, fila))
        _f(ws, 'E%d' % fila, '=IF(ROUND(SUM(B%d:D%d),4)=1,"OK","Adjust: the three shares must add up to 100%%")'
           % (fila, fila))
    f_mix1 = fila
    rng_paq = '$A$%d:$A$%d' % (f_mix0, f_mix1)
    fila += 2
    _f(ws, 'A%d' % fila, 'ASSUMPTIONS (Kit estimate — edit)', b=True, fill=GRIS)
    for col, t in zip('BE', ('Value', 'Notes')):
        _f(ws, '%s%d' % (col, fila), t, b=True, fill=GRIS)
    for k, rot, fmt, val, nota in COMPRA:
        fila += 1
        _f(ws, 'A%d' % fila, rot)
        _f(ws, 'B%d' % fila, val, fmt=fmt if fmt != '@' else None, fill=VERDE, locked=False)
        if nota:
            _f(ws, 'E%d' % fila, nota, wrap=True)
        R[k] = '$B$%d' % fila
    fila += 2
    _f(ws, 'A%d' % fila, 'SPIRITS SPLIT (share of the cocktails & spirits; Kit estimate — edit)', b=True, fill=GRIS)
    _f(ws, 'B%d' % fila, 'Share', b=True, fill=GRIS)
    lic0 = fila + 1
    for n, s in LICORES.items():
        fila += 1
        _f(ws, 'A%d' % fila, n)
        _f(ws, 'B%d' % fila, s, fmt=PCT, fill=VERDE, locked=False)
        pct_celdas.append('B%d' % fila)
    lic1 = fila
    fila += 1
    _f(ws, 'A%d' % fila, 'Check (must be 100%)')
    _f(ws, 'B%d' % fila, '=IF(ROUND(SUM(B%d:B%d),4)=1,"OK","Adjust to 100%%")' % (lic0, lic1))
    # cálculo
    def share(col):
        return 'INDEX($%s$%d:$%s$%d,MATCH(%s,%s,0))' % (col, f_mix0, col, f_mix1, R['package'], rng_paq)
    up = '(1+%s)' % R['buffer']
    fila += 2
    _f(ws, 'A%d' % fila, 'DRINKS BY CATEGORY', b=True, fill=GRIS)
    _f(ws, 'B%d' % fila, 'Drinks', b=True, fill=GRIS)
    for k, rot, col in (('d_cocktails', 'Cocktails & spirits', 'B'), ('d_wine', 'Wine', 'C'), ('d_beer', 'Beer', 'D')):
        fila += 1
        _f(ws, 'A%d' % fila, rot)
        _f(ws, 'B%d' % fila, '=IFERROR(%s*%s,0)' % (R['drinks'], share(col)), fmt=DEC1)
        R[k] = '$B$%d' % fila
    fila += 2
    quien = '=IF(%s=%s,"You (licensed package)","The host (dry hire)")' % (R['package'], T_LICENCIA)
    _f(ws, 'A%d' % fila, 'SHOPPING LIST — ALCOHOL', b=True, fill=GRIS)
    _f(ws, 'B%d' % fila, 'Quantity', b=True, fill=GRIS)
    _f(ws, 'C%d' % fila, 'Unit', b=True, fill=GRIS)
    _f(ws, 'D%d' % fila, quien, b=True, fill=GRIS, wrap=True)
    _f(ws, 'E%d' % fila, 'Notes', b=True, fill=GRIS)
    ws.row_dimensions[fila].height = 30
    filas_compra = OrderedDict()
    fila += 1
    _f(ws, 'A%d' % fila, 'Spirits, total (750 ml bottles)', b=True)
    _f(ws, 'B%d' % fila, '=ROUNDUP(%s*%s/%s,0)' % (R['d_cocktails'], up, R['pours']), fmt=ENTERO, b=True)
    _f(ws, 'C%d' % fila, 'bottles')
    _f(ws, 'E%d' % fila, 'Cocktails & spirits ÷ drinks per bottle, with the buffer.', wrap=True)
    filas_compra['b_spirits'] = fila
    for i, (n, _s) in enumerate(LICORES.items()):
        fila += 1
        _f(ws, 'A%d' % fila, '   %s' % n)
        _f(ws, 'B%d' % fila, '=ROUNDUP(%s*$B$%d*%s/%s,0)' % (R['d_cocktails'], lic0 + i, up, R['pours']), fmt=ENTERO)
        _f(ws, 'C%d' % fila, 'bottles')
        filas_compra['b_lic_%d' % i] = fila
    _f(ws, 'E%d' % (filas_compra['b_lic_0']), 'Rounded up type by type, so the split can add up to a little more '
                                               'than the total.', wrap=True)
    for k, rot, fx, unidad, nota in (
            ('b_wine', 'Wine (750 ml bottles)', '=ROUNDUP(%s*%s/%s,0)' % (R['d_wine'], up, R['wine_glasses']), 'bottles',
             'Split between red, white and sparkling to taste.'),
            ('u_beer', 'Beer (12 oz cans or bottles)', '=ROUNDUP(%s*%s*%s,0)' % (R['d_beer'], R['beer_per'], up), 'units',
             ''),
            ('b_toast', 'Sparkling wine for the toast (750 ml bottles)',
             '=IF(%s="Yes",ROUNDUP(%s/%s,0),0)' % (R['toast'], R['guests'], R['toast_pours']), 'bottles',
             'Only if the switch above is Yes.')):
        fila += 1
        _f(ws, 'A%d' % fila, rot)
        _f(ws, 'B%d' % fila, fx, fmt=ENTERO)
        _f(ws, 'C%d' % fila, unidad)
        if nota:
            _f(ws, 'E%d' % fila, nota, wrap=True)
        filas_compra[k] = fila
    fila += 2
    _f(ws, 'A%d' % fila, 'WHAT THE BAR BRINGS', b=True, fill=GRIS)
    for col, t in zip('BCE', ('Quantity', 'Unit', 'Notes')):
        _f(ws, '%s%d' % (col, fila), t, b=True, fill=GRIS)
    for k, rot, fx, fmt, unidad, nota in (
            ('ice_lb', 'Ice', '=%s*%s' % (R['guests'], g + '$B$%d' % K['ice_lb']), DEC1, 'lb',
             'Commercial bagged ice. Ask your county about ice and garnish rules.'),
            ('ice_bags', 'Ice, in bags', '=ROUNDUP(B{prev}/%s,0)' % R['bag_lb'], ENTERO, 'bags', ''),
            ('cups_n', 'Cups', '=ROUNDUP(%s*%s*%s,0)' % (R['drinks'], R['cups_per'], up), ENTERO, 'units', ''),
            ('napkins_n', 'Napkins', '=ROUNDUP(%s*%s*%s,0)' % (R['drinks'], R['napkins_per'], up), ENTERO, 'units', ''),
            ('mixer_servings', 'Mixer and garnish servings', '=ROUNDUP(%s*%s,0)' % (R['d_cocktails'], up), ENTERO,
             'servings', 'One per cocktail: see the "%s" tab for each recipe.' % H3)):
        fila += 1
        _f(ws, 'A%d' % fila, rot)
        _f(ws, 'B%d' % fila, fx.replace('{prev}', str(fila - 1)), fmt=fmt)
        _f(ws, 'C%d' % fila, unidad)
        if nota:
            _f(ws, 'E%d' % fila, nota, wrap=True)
        filas_compra[k] = fila
    _dv(ws, 'decimal', pct_celdas, f1='0', f2='1', op='between', titulo='Invalid percentage',
        msg='Enter a share between 0 and 1 (0.25 = 25%).')
    _dv(ws, 'whole', [R[k].replace('$', '') for k in ('pours', 'wine_glasses', 'toast_pours', 'bag_lb')], f1='1',
        f2='1000', op='between', titulo='Invalid count', msg='Enter a whole number of 1 or more.')
    _dv(ws, 'decimal', [R[k].replace('$', '') for k in ('beer_per', 'cups_per', 'napkins_per')], f1='0',
        op='greaterThanOrEqual', titulo='Invalid amount', msg='Enter a number greater than or equal to 0.')
    _dv(ws, 'decimal', [R['buffer'].replace('$', '')], f1='0', f2='1', op='between', titulo='Invalid percentage',
        msg='Enter a percentage between 0 and 1 (0.10 = 10%).')
    _dv(ws, 'list', [R['toast'].replace('$', '')], f1='"Yes,No"', titulo='Invalid entry', msg='Choose Yes or No.')
    ws.freeze_panes = 'A4'
    _imprimir(ws, horizontal=True)
    return ws, filas_compra


def hoja_specs(wb):
    ws = wb.create_sheet(H3)
    anchos = (('A', 34), ('B', 11), ('C', 8), ('D', 12), ('E', 11), ('F', 9), ('G', 11), ('H', 12), ('I', 11),
              ('J', 11), ('K', 11), ('L', 12))
    for col, w in anchos:
        ws.column_dimensions[col].width = w
    _cabecera(ws, 'SIGNATURE COCKTAIL SPECS — 10 RECIPES, SCALED AND COSTED',
              'Ten signature cocktails, including one zero-proof mocktail, in US ounces. Type your local price for each '
              'bottle or pack in the green cells: the cost per drink and the batch quantities for 50, 100 and 200 drinks '
              'are calculated. In dry hire, everything that contains alcohol (bitters included) goes on the host\'s '
              'shopping list; you supply mixers, juices, syrups, garnishes and ice. Pack prices are kit estimates.', 'L')
    cab = ('Ingredient', 'Per drink', 'Unit', 'Who buys (dry hire)', 'Pack size', 'Pack unit', 'Pack price',
           'Cost per drink', '50 drinks', '100 drinks', '200 drinks', 'Packs for 100 drinks')
    fila = 4
    for col, t in zip('ABCDEFGHIJKL', cab):
        _f(ws, '%s%d' % (col, fila), t, b=True, fill=GRIS, wrap=True)
    ws.row_dimensions[fila].height = 30
    verdes_dec = []
    filas = []
    for nombre, vaso, alergeno, ings in RECETAS:
        fila += 2
        _f(ws, 'A%d' % fila, nombre.upper(), b=True, fill=GRIS)
        _f(ws, 'B%d' % fila, vaso + ('. ' + alergeno if alergeno else ''), i=True, fill=GRIS)
        ws.merge_cells('B%d:L%d' % (fila, fila))
        f0 = fila + 1
        for ing, cant, u, quien, env, ue, precio in ings:
            fila += 1
            _f(ws, 'A%d' % fila, ing)
            _f(ws, 'B%d' % fila, cant, fmt='0.00', fill=VERDE, locked=False)
            _f(ws, 'C%d' % fila, u)
            _f(ws, 'D%d' % fila, quien)
            _f(ws, 'E%d' % fila, env, fmt='0.0', fill=VERDE, locked=False)
            _f(ws, 'F%d' % fila, ue)
            _f(ws, 'G%d' % fila, precio, fmt=MONEDA, fill=VERDE, locked=False)
            _f(ws, 'H%d' % fila, '=IF(E%d=0,0,B%d/E%d*G%d)' % (fila, fila, fila, fila), fmt='0.000')
            for col, n in zip('IJK', ESCALAS):
                _f(ws, '%s%d' % (col, fila), '=B%d*%d' % (fila, n), fmt='#,##0.00')
            _f(ws, 'L%d' % fila, '=IF(E%d=0,0,ROUNDUP(B%d*100/E%d,0))' % (fila, fila, fila), fmt=ENTERO)
            verdes_dec += ['B%d' % fila, 'G%d' % fila]
        f1 = fila
        fila += 1
        _f(ws, 'A%d' % fila, 'Cost per drink, all ingredients', b=True)
        _f(ws, 'H%d' % fila, '=SUM(H%d:H%d)' % (f0, f1), fmt=MONEDA, b=True)
        fila += 1
        _f(ws, 'A%d' % fila, 'Your cost per drink in dry hire (bar-supplied items)', b=True)
        _f(ws, 'H%d' % fila, '=SUMIF(D%d:D%d,"%s",H%d:H%d)' % (f0, f1, _B, f0, f1), fmt=MONEDA, b=True)
        filas.append((f0, f1, fila - 1, fila))
        _dv(ws, 'decimal', ['E%d' % r for r in range(f0, f1 + 1)], f1='0.01', op='greaterThanOrEqual',
            titulo='Invalid pack size', msg='Enter the pack size in the pack unit (greater than 0).')
    _dv(ws, 'decimal', verdes_dec, f1='0', op='greaterThanOrEqual', titulo='Invalid amount',
        msg='Enter a number greater than or equal to 0.')
    fila += 2
    for t in ('Specs converted from milliliters (30 ml = 1 oz) and rounded to the nearest 0.25 oz. A 750 ml bottle '
              'holds about 25.4 oz and a 1 liter bottle about 33.8 oz.',
              'Batching tip: pre-batch the still part of each cocktail (spirits, juice, syrup) and add sparkling wine '
              'or soda at the bar. Always offer at least one zero-proof drink.',
              'Allergens: tell guests that the Mai Tai contains almond (orgeat) and that the Pina Colada and the '
              'mocktail contain coconut. Keep an allergen card at the bar.'):
        _f(ws, 'A%d' % fila, t, wrap=True)
        ws.merge_cells('A%d:L%d' % (fila, fila))
        ws.row_dimensions[fila].height = 30
        fila += 1
    ws.freeze_panes = 'B5'
    _imprimir(ws, horizontal=True)
    return ws, filas


def hoja_ejemplos(wb):
    ws = wb.create_sheet(H4)
    ws.column_dimensions['A'].width = 60
    for col in 'BCD':
        ws.column_dimensions[col].width = 26
    _cabecera(ws, '3 EXAMPLE QUOTES — calculated with the same formulas as the "%s" tab' % H1,
              'Kit estimates for illustration only, not market prices: change the inputs in "%s" to price your own '
              'events. All three are dry hire: the host buys the alcohol.' % H1, 'D')
    _f(ws, 'A4', 'INPUTS', b=True, fill=GRIS)
    for j, nombre in enumerate(EJEMPLOS):
        _f(ws, '%s4' % 'BCD'[j], nombre, b=True, fill=GRIS, wrap=True)
    ws.row_dimensions[4].height = 30
    for k, rot, fmt, val, nota in ENTRADAS:
        _f(ws, 'A%d' % K[k], rot)
        for j, e in enumerate(EJEMPLOS.values()):
            _f(ws, '%s%d' % ('BCD'[j], K[k]), e[k], fmt=fmt if fmt != '@' else None, wrap=k == 'package')
        if k == 'package':
            ws.row_dimensions[K[k]].height = 30
    _f(ws, 'A%d' % F_CALC, 'RESULT', b=True, fill=GRIS)
    for j in range(3):
        _f(ws, '%s%d' % ('BCD'[j], F_CALC), list(EJEMPLOS)[j], b=True, fill=GRIS, wrap=True)
    ws.row_dimensions[F_CALC].height = 30
    for k, rot, fmt, expr, nota in SALIDAS:
        neg = k in ('price', 'total', 'margin_calc')
        _f(ws, 'A%d' % S[k], rot, b=neg)
        for j in range(3):
            col = 'BCD'[j]
            ref = {kk: '$%s$%d' % (col, f) for kk, f in K.items()}
            ref.update({kk: '$%s$%d' % (col, f) for kk, f in S.items()})
            _f(ws, '%s%d' % (col, S[k]), formula_salida(k, expr, ref), fmt=fmt, b=neg)
    f = S[SALIDAS[-1][0]] + 2
    _f(ws, 'A%d' % f, 'WHAT EACH EXAMPLE SHOWS', b=True, fill=GRIS)
    for j, (nombre, nota) in enumerate(EJ_NOTAS.items()):
        _f(ws, 'A%d' % (f + 1 + j), nombre, b=True)
        _f(ws, 'B%d' % (f + 1 + j), nota, wrap=True)
        ws.merge_cells('B%d:D%d' % (f + 1 + j, f + 1 + j))
        ws.row_dimensions[f + 1 + j].height = 45
    ws.freeze_panes = 'B5'
    _imprimir(ws, horizontal=True)
    return ws


INSTRUCCIONES = [
    ('HOW TO USE THE MOBILE BAR PRICING CALCULATOR', True),
    ('This workbook turns your costs into a price per guest and a full quote, and builds the shopping list for the '
     'event. You type only in the GREEN cells; the gross margin is calculated, never typed. The default model is dry '
     'hire ' + GLOSA + '.', False),
    ('1. "%s": choose the bar package, then enter guests, hours of service and the drinks rule (2 drinks per guest in '
     'the first hour and 1 per hour after: eventplanning.com). The suggested number of bartenders comes from the '
     'table at the bottom of the tab (guests per bartender by package: edit it to your own standard).' % H1, False),
    ('2. Enter the staff (bartenders, barbacks, paid hours, wages, payroll taxes) and the costs you supply: mixers and '
     'garnishes per drink, cups and napkins per guest, ice and mileage. The alcohol cost only counts with the licensed '
     'package.', False),
    ('3. Set a flat bar fee if you charge one, the specialty bar upcharge, the marketplace fee, your target margin, the '
     'service charge, your sales tax rate and the switch "Does your state tax the bar service and the service '
     'charge?". The price per guest is the direct cost divided by (1 - target margin - marketplace fee), minus the '
     'flat fee, rounded up to the next 0.50; the subtotal never goes below your minimum.', False),
    ('4. "%s": the drink mix of each package and the bottle math give the host\'s shopping list (spirits split by '
     'type, wine, beer, sparkling wine for the toast) and what the bar brings (ice in bags, cups, napkins, mixer '
     'servings). Send it to the host about two weeks before the event.' % H2, False),
    ('5. "%s": ten recipes with the amount per drink, who buys each ingredient in dry hire, the cost per drink from the '
     'pack prices you type, and the quantities for 50, 100 and 200 drinks.' % H3, False),
    ('6. "%s": three complete quotes calculated with the same formulas.' % H4, False),
    ('DRY HIRE: WHAT YOU CAN AND CANNOT CHARGE', True),
    ('In dry hire the host buys and owns the alcohol and keeps what is left. Do not sell drinks, drink tickets or '
     'admission that includes drinks, do not add a markup on the alcohol, and keep any tip jar only where your state '
     'allows it: Missouri treats admission fees, tips and alcohol surcharges as possible sales (Missouri Division of '
     'Alcohol and Tobacco Control, Industry Circular 24-01). Rules vary by state: check yours before you quote.',
     False),
    ('SERVICE CHARGE, GRATUITY AND SALES TAX (US)', True),
    ('A mandatory service charge is not a tip: the part you pay out to the staff is wages, with payroll taxes (IRS Rev. '
     'Rul. 2012-18), and the part you keep is revenue. A voluntary gratuity is the client\'s decision and belongs to the '
     'staff: this calculator keeps it out of the price and the total.', False),
    ('Sales tax: California taxes a caterer\'s service even when the client supplies the drinks (cdtfa.ca.gov), so '
     'a bartending service is taxable there; other states differ. Set your rate and the switch.',
     False),
    ('UK: if you are VAT-registered, enter 0.20 as the tax rate. Selling alcohol needs a premises licence or a '
     'Temporary Event Notice: see section 9 of the plan.', False),
    ('KIT ESTIMATES', True),
    ('Every default value, every pack price and the three example quotes are kit estimates for illustration, not '
     'market prices. Replace them with your own wages, supplier prices and tax rate before you quote.', False),
]


def hoja_instrucciones(wb):
    ws = wb.create_sheet(H5)
    ws.column_dimensions['A'].width = 110
    f = 1
    for t, b in INSTRUCCIONES:
        _f(ws, 'A%d' % f, t, b=b, sz=13 if f == 1 else 11, wrap=True)
        if not b:
            ws.row_dimensions[f].height = max(16, 15 * math.ceil(len(t) / 120.0))
        f += 2 if b and f > 1 else 1
    f += 1
    for t in ('To edit a cell that is not green: Review → Unprotect Sheet (there is no password).',
              'More templates and kits in the catalog at ' + M.URL_TIENDA, M.FIRMA_EN, M.version_en('mbar'),
              M.AVISO_EN):
        _f(ws, 'A%d' % f, t, wrap=True, i=t == M.AVISO_EN)
        f += 1
    _imprimir(ws)
    return ws


def docprops(wb):
    p = M.PLANES['mbar']
    pr = wb.properties
    pr.title = '%s · %s' % (TITULO, p['producto'])
    pr.subject = '%s · v%s' % (p['producto'], M.VERSION)
    pr.keywords = p['keywords']
    pr.description = p['url']
    pr.category = 'AI Chef Pro · Digital products'
    pr.creator = 'AI Chef Pro'
    pr.lastModifiedBy = 'AI Chef Pro'


def generar(destino):
    wb = openpyxl.Workbook()
    hoja_calculadora(wb)
    hoja_compra(wb)
    hoja_specs(wb)
    hoja_ejemplos(wb)
    hoja_instrucciones(wb)
    docprops(wb)
    os.makedirs(os.path.dirname(destino), exist_ok=True)
    wb.save(destino)
    sys.path.insert(0, PD)
    import inject_cache as IC                                                    # noqa: E402
    nf, inj, fail = IC.inject(destino)
    print('%s: formulas=%d cache_inyectado=%d fallos_pycel=%d' % (os.path.relpath(destino, REPO), nf, inj, fail))
    return fail


# --------------------------------------------------------------------------
# Autotest (D90): pycel = caché = Python en las 4 pestañas de cálculo + el paquete con licencia en la hoja principal
# --------------------------------------------------------------------------
def _cmp(err, etiqueta, got, esperado, tol=0.005):
    if not isinstance(got, (int, float)) or abs(float(got) - esperado) > tol:
        err.append('%s: %r ≠ esperado %r' % (etiqueta, got, esperado))


def autotest(destino):
    from pycel import ExcelCompiler
    err, n = [], 0
    xl = ExcelCompiler(filename=destino)
    wv = openpyxl.load_workbook(destino, data_only=True)
    wf = openpyxl.load_workbook(destino)
    defecto = OrderedDict((k, v) for k, _r, _f2, v, _n in ENTRADAS)
    # 1) hoja principal y los 3 ejemplos (pycel + caché)
    casos = [(H4, 'BCD'[j], e, nombre) for j, (nombre, e) in enumerate(EJEMPLOS.items())]
    casos.append((H1, 'B', defecto, 'hoja principal (valores por defecto)'))
    for hoja, col, e, nombre in casos:
        esperado = calcular(e)
        for k, v in esperado.items():
            c = '%s%d' % (col, S[k])
            n += 1
            _cmp(err, '%s · %s!%s (%s) pycel' % (nombre, hoja, c, k), xl.evaluate("'%s'!%s" % (hoja, c)), v)
            _cmp(err, '%s · %s!%s (%s) caché' % (nombre, hoja, c, k), wv[hoja][c].value, v)
        if esperado['margin_calc'] + 1e-9 < e['margin']:
            err.append('%s: margen calculado %.4f < objetivo %.4f' % (nombre, esperado['margin_calc'], e['margin']))
        if e['package'] != FULL_LIC and esperado['c_alcohol'] != 0:
            err.append('%s: alcohol en el coste de un paquete dry hire' % nombre)
    # 2) Drinks & Shopping List con los valores por defecto (pycel + caché)
    _ws2, filas_compra = None, None
    wb_tmp = openpyxl.Workbook()
    _ws2, filas_compra = hoja_compra(wb_tmp)                                   # mismas filas que el fichero
    compra = calcular_compra(defecto)
    for k, v in compra.items():
        if k not in filas_compra and not k.startswith('d_'):
            continue
        if k.startswith('d_'):
            continue                                                         # intermedias: se comprueban vía botellas
        c = 'B%d' % filas_compra[k]
        n += 1
        _cmp(err, 'compra %s!%s (%s) pycel' % (H2, c, k), xl.evaluate("'%s'!%s" % (H2, c)), v)
        _cmp(err, 'compra %s!%s (%s) caché' % (H2, c, k), wv[H2][c].value, v)
    # 3) Signature Cocktail Specs (pycel + caché): coste por ingrediente, escalado, envases y totales
    wb_tmp2 = openpyxl.Workbook()
    _ws3, filas_specs = hoja_specs(wb_tmp2)
    for (nombre, filas, tot, bar), (f0, f1, ft, fb) in zip(calcular_specs(), filas_specs):
        for i, x in enumerate(filas):
            r = f0 + i
            for col, k in (('H', 'coste'), ('I', 'x50'), ('J', 'x100'), ('K', 'x200'), ('L', 'envases100')):
                n += 1
                _cmp(err, 'specs %s %s%d pycel' % (nombre, col, r), xl.evaluate("'%s'!%s%d" % (H3, col, r)), x[k], 1e-6)
                _cmp(err, 'specs %s %s%d caché' % (nombre, col, r), wv[H3]['%s%d' % (col, r)].value, x[k], 1e-6)
        for fila, v in ((ft, tot), (fb, bar)):
            n += 1
            _cmp(err, 'specs %s H%d pycel' % (nombre, fila), xl.evaluate("'%s'!H%d" % (H3, fila)), v, 1e-6)
            _cmp(err, 'specs %s H%d caché' % (nombre, fila), wv[H3]['H%d' % fila].value, v, 1e-6)
        if bar > tot + 1e-9 or bar <= 0:
            err.append('specs %s: coste del bar %.3f fuera de (0, total %.3f]' % (nombre, bar, tot))
    # 4) el paquete con licencia en la hoja principal (solo pycel: la caché es la del paquete por defecto)
    lic = dict(defecto, package=FULL_LIC)
    xl.set_value("'%s'!B%d" % (H1, K['package']), FULL_LIC)
    esperado = calcular(lic)
    for k, v in esperado.items():
        n += 1
        _cmp(err, 'licencia · %s!B%d (%s) pycel' % (H1, S[k], k), xl.evaluate("'%s'!B%d" % (H1, S[k])), v)
    if not esperado['c_alcohol'] > 0 or not esperado['price'] > calcular(defecto)['price']:
        err.append('paquete con licencia sin efecto (alcohol %r)' % esperado['c_alcohol'])
    compra_lic = calcular_compra(lic)
    for k in ('b_spirits', 'b_wine', 'u_beer'):
        n += 1
        _cmp(err, 'licencia · compra %s pycel' % k, xl.evaluate("'%s'!B%d" % (H2, filas_compra[k])), compra_lic[k])
    # 5) el interruptor del sales tax cambia el resultado
    e = dict(EJEMPLOS['Wedding, 120 guests, 5 hours'])
    if not calcular(dict(e, tax_sc='Yes'))['tax_amt'] > calcular(dict(e, tax_sc='No'))['tax_amt']:
        err.append('interruptor Yes/No sin efecto')
    # higiene del fichero (censo-entregables + D20 + D8)
    p = wf.properties
    if p.creator != 'AI Chef Pro' or not (p.title or '').startswith(TITULO) or M.PLANES['mbar']['url'] != p.description:
        err.append('docProps D20: %r / %r / %r' % (p.creator, p.title, p.description))
    textos = []
    for ws in wf.worksheets:
        if ws.page_setup.paperSize != 1 or not ws.sheet_properties.pageSetUpPr.fitToPage or \
                ws.oddFooter.center.text != M.FOOTER_EN:
            err.append('%s: impresión (Letter + ajuste + pie)' % ws.title)
        if not ws.protection.sheet:
            err.append('%s: hoja sin proteger' % ws.title)
        for row in ws.iter_rows():
            for c in row:
                v = c.value
                if v == '':
                    err.append('%s!%s: cadena vacía' % (ws.title, c.coordinate))
                if isinstance(v, str) and c.data_type != 'f':
                    textos.append((ws.title, c.coordinate, v))
                if isinstance(v, str) and c.data_type == 'f':
                    textos += [(ws.title, c.coordinate + ' (fórmula)', lit) for lit in re.findall(r'"((?:[^"]|"")*)"', v)
                               if any(ch.isalpha() for ch in lit)]
                if c.data_type == 'f' and wv[ws.title][c.coordinate].value is None:
                    err.append('%s!%s: fórmula sin caché' % (ws.title, c.coordinate))
                if c.fill.fgColor.rgb in ('00E8F5E9', 'FFE8F5E9') and c.protection.locked:
                    err.append('%s!%s: verde bloqueada' % (ws.title, c.coordinate))
                if isinstance(v, str) and c.data_type != 'f' and '€' in v:
                    err.append('%s!%s: «€»' % (ws.title, c.coordinate))
                fmt = c.number_format or ''
                if '€' in fmt or '$' in fmt:
                    err.append('%s!%s: símbolo de moneda en el formato (D8)' % (ws.title, c.coordinate))
        for dv in ws.data_validations.dataValidation:
            for t in (dv.error, dv.errorTitle, dv.prompt, dv.promptTitle):
                if t:
                    textos.append((ws.title, 'DV %s' % dv.sqref, t))
    for campo in ('title', 'subject', 'keywords', 'description'):
        textos.append(('docProps', campo, getattr(p, campo) or ''))
    for h, c, t in textos:
        r = [x for x in M.restos_espanol(t) if not x[0].startswith('oficio')]
        if r:
            err.append('%s!%s: restos de español %s' % (h, c, r[:1]))
        if M.no_latinos(t):
            err.append('%s!%s: caracteres no latinos' % (h, c))
        if M.cifras_prohibidas(t):
            err.append('%s!%s: cifra/término prohibido %s' % (h, c, M.cifras_prohibidas(t)[:1]))
        if re.search(r'(?i)\bvalue\b.*\$\d|\bsave \$|total value|\bwas \$|\d+\s?% off\b', t):
            err.append('%s!%s: ancla de precio' % (h, c))
    hojas = [ws.title for ws in wf.worksheets]
    if hojas != [H1, H2, H3, H4, H5]:
        err.append('pestañas %r' % hojas)
    if not any(GLOSA in t for _h, _c, t in textos if _h == H1):
        err.append('falta la glosa de «dry hire» en la primera mención (%s)' % H1)
    for x in err:
        print('ERROR', x)
    print('autotest calculadora: %d comprobaciones numéricas (pycel + caché) en %s + el paquete con licencia · %d '
          'textos · %s' % (n, ', '.join(HOJAS_CALCULO), len(textos), 'VERDE' if not err else 'ROJO (%d errores)' % len(err)))
    for nombre, e in list(EJEMPLOS.items()) + [('Main tab (defaults)', defecto), ('Main tab, licensed', lic)]:
        r = calcular(e)
        print('  %-45s %3d inv · %4.0f copas · coste %8.2f · precio/inv %6.2f · subtotal %8.2f · /h %7.2f · total %8.2f · '
              'margen %.1f%% · bartenders sug. %d' % (nombre, e['guests'], r['drinks'], r['c_total'], r['price'],
                                                     r['subtotal'], r['per_hour_eq'], r['total'], r['margin_calc'] * 100,
                                                     r['bt_sug']))
    c = calcular_compra(defecto)
    print('  compra (defecto): %d botellas de destilado %s · vino %d · cerveza %d · hielo %.0f lb = %d bolsas' % (
        c['b_spirits'], [c['b_lic_%d' % i] for i in range(len(LICORES))], c['b_wine'], c['u_beer'], c['ice_lb'],
        c['ice_bags']))
    for nombre, _f2, tot, bar in calcular_specs():
        print('  %-32s coste/copa %.2f · del bar %.2f' % (nombre, tot, bar))
    return 1 if err else 0


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--autotest', action='store_true', help='solo el autotest sobre el fichero existente')
    ap.add_argument('--salida', default=DESTINO)
    a = ap.parse_args()
    if not a.autotest:
        if generar(a.salida):
            print('ERROR: fallos de pycel en inject_cache')
            return 1
    return autotest(a.salida)


if __name__ == '__main__':
    sys.exit(main())
