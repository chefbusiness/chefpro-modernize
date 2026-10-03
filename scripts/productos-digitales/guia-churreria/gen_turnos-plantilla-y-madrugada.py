#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
gen_turnos-plantilla-y-madrugada.py - libro 8 de «Cómo Montar una
Churrería-Chocolatería» (producto 51, `guia-churreria-chocolateria`; SPEC
`scripts/productos-digitales/guia-churreria-SPEC.md`, §2.1 fila 8, §2.2, §3.2 y §5).

Molde N (SPEC §2.3): sin generador hermano que calcar. Se toman los helpers de
`guias-v2_0/motor.py`, la lógica de la hoja Personal del plan financiero de la hermana
(coste por puesto = MAX(salario; SMI) más la Seguridad Social de empresa) y el cierre
de los libros 3 y 4 de este mismo producto (inject_cache + verificación data_only +
contrato de cruces + lista negra + mapa + demo con pycel).

Hojas (`datos_ejemplo.HOJAS[8]`): Instrucciones · Parámetros · Cuadrante por Franja ·
Nocturnidad ET y Convenio · Coste de Plantilla · Identifica tu Convenio · Horas del
Titular.

LAS DOS FIGURAS (A1, D9), QUE NUNCA SE MEZCLAN
----------------------------------------------
  · Trabajador nocturno del ET art. 36.1 (CUN-29): horas entre las 22:00 y las 6:00;
    tres horas diarias o un tercio de la jornada anual. El libro da «sí», «no» o
    «revísalo con tu asesor» (el «normalmente» del artículo no está interpretado).
  · Plus de nocturnidad del convenio de ejemplo (CUN-40): +1 % de 22:00 a 0:00 y +25 %
    de 0:00 a 8:00, con las dos franjas en celda verde, y el umbral de cinco horas.
El coste anual por puesto es MAX(salario; SMI anual a su jornada) con semáforo (A7).
Convenio de Madrid de EJEMPLO, tablas 2025, vencido y en negociación (V-01, rama A).

CRUCES (D16, `datos_ejemplo.CRUCES`)
------------------------------------
Recibe X6 (personas en la línea en el pico, del libro 1), X7 (horas por franja de la
semana tipo y horas de refuerzo, del libro 5) y X8 (coste hora del escandallo, del
libro 3): celdas VERDES de «Parámetros» con el rótulo «cópialo del libro N: ...» y su
fila de CUADRE con semáforo. Nunca una fórmula entre ficheros. X6 alimenta además el
mínimo de personas en la línea del pico en «Cuadrante por Franja», de donde salen los
huecos de cobertura. Es ORIGEN de X15 (coste anual de plantilla, `Coste de
Plantilla`!L5, al libro 6).

FRONTERA (SR-14)
----------------
No reproduce el cuadrante semana a semana por empleado ni las nóminas del Kit Gestión
de Personal y Turnos: trabaja con la SEMANA TIPO que usa el plan (horas por franja, las
dos figuras de nocturnidad, el plus y el coste anual por puesto). Venta cruzada: «para
el cuadrante semana a semana, el Kit Gestión de Personal y Turnos».

Salida: build/turnos-plantilla-y-madrugada.xlsx + build/mapa-...json.
Uso: /usr/local/bin/python3 gen_turnos-plantilla-y-madrugada.py
Via: Claude Code
"""
import importlib.util
import json
import os
import re
import subprocess
import sys

import openpyxl
from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter, column_index_from_string

AQUI = os.path.dirname(os.path.abspath(__file__))
RAIZ = os.path.normpath(os.path.join(AQUI, '..'))
if AQUI not in sys.path:
    sys.path.insert(0, AQUI)

import _comun_churreria as CC                                  # noqa: E402
import motor                                                   # noqa: E402


def _importar_datos():
    spec = importlib.util.spec_from_file_location(
        'datos_ejemplo_churreria', os.path.join(AQUI, 'datos_ejemplo.py'))
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    return mod


D = _importar_datos()
import enmiendas_churreria                                     # noqa: E402
enmiendas_churreria.aplicar(D)      # R-04/05/06/13/14/15: el contrato de cruces enmendado
motor.CTX['producto'] = CC.PID

LIBRO = 8
FICHERO = D.LIBROS[LIBRO]
NOMBRE = os.path.splitext(FICHERO)[0]
TITULO = 'Turnos, Plantilla y Madrugada'
SUBJECT = CC.PRODUCTO + ' · Versión 1.0 · octubre 2026'

(H_INS, H_PAR, H_CUA, H_NOC, H_COS, H_IDE, H_TIT) = D.HOJAS[LIBRO]
assert NOMBRE == 'turnos-plantilla-y-madrugada'


def Q(hoja):
    return "'" + hoja + "'!"


def ABS_(hoja, coord):
    """Referencia absoluta a otra hoja del MISMO libro."""
    m = re.match(r'([A-Z]+)(\d+)$', coord)
    return Q(hoja) + '$' + m.group(1) + '$' + m.group(2)


def RNG(hoja, c1, c2):
    m1 = re.match(r'([A-Z]+)(\d+)$', c1)
    m2 = re.match(r'([A-Z]+)(\d+)$', c2)
    return '%s$%s$%s:$%s$%s' % (Q(hoja), m1.group(1), m1.group(2), m2.group(1), m2.group(2))


def col(letra, n):
    """Letra de columna desplazada n posiciones."""
    return get_column_letter(column_index_from_string(letra) + n)


# --------------------------------------------------------------------------
# Formato y utilidades (segunda capa; la primera es `_comun_churreria.py`)
# --------------------------------------------------------------------------
ORO = 'FFD700'
GRIS = '888888'
CAB_BG, CAB_FG = '2D2D2D', 'FFFFFF'
CREMA = 'FFF8E1'
EUR = motor.FMT_EUR
PCT = motor.FMT_PCT
ENT = motor.FMT_ENT
DEC = '#,##0.00'
DEC1 = '#,##0.0'
DEC4 = '#,##0.0000'
HORA = 'h:mm'

NOTAS_LEGALES = []
VERDES = []
CRUCE_VERDES = []        # (hoja, coord, x, concepto)
CRUCE_CUADRES = []
CONTRASTES = []       # (hoja, coord, x, concepto)
SI, NO = 'sí', 'no'
SI_NO = (SI, NO)


def cabecera_hoja(ws, titulo, nota_txt=None):
    motor.val(ws, 'A1', CC.titulo_hoja(titulo), bold=True)
    ws['A1'].font = Font(bold=True, size=16, color=ORO)
    motor.val(ws, 'A2', CC.SUBTITULO)
    ws['A2'].font = Font(size=10, color=GRIS)
    if nota_txt:
        motor.val(ws, 'A3', nota_txt)
        ws['A3'].font = Font(italic=True, size=9)


def seccion(ws, coord, texto):
    motor.val(ws, coord, texto, bold=True)
    ws[coord].font = Font(bold=True, size=12, color=ORO)


def gris(ws, coord, texto, alto=None):
    motor.val(ws, coord, texto, wrap=True)
    ws[coord].font = Font(size=8, color=GRIS)
    if alto:
        ws.row_dimensions[int(re.sub(r'[A-Z]', '', coord))].height = alto
    return ws[coord]


def encabezados(ws, fila, cols, alto=42, congelar=True):
    """`cols` = [(letra, texto, ancho|None), ...]."""
    for letra, txt, ancho in cols:
        cel = ws[letra + str(fila)]
        cel.value = txt
        cel.font = Font(bold=True, color=CAB_FG, size=9)
        cel.fill = PatternFill('solid', fgColor=CAB_BG)
        cel.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        if ancho is not None:
            ws.column_dimensions[letra].width = ancho
    ws.row_dimensions[fila].height = alto
    if congelar:
        ws.freeze_panes = 'A%d' % (fila + 1)


def crema(cel):
    cel.fill = PatternFill('solid', fgColor=CREMA)
    return cel


def entrada(ws, coord, valor, fmt=None, etiqueta=''):
    """Celda VERDE con su valor por defecto. Ninguna verde vacía."""
    if valor is None or (isinstance(valor, str) and not valor.strip()):
        raise SystemExit('entrada(%s!%s, «%s») sin valor por defecto: ninguna celda '
                         'verde puede quedar vacía' % (ws.title, coord, etiqueta))
    cel = motor.val(ws, coord, valor, fmt=fmt, verde_=True)
    VERDES.append((ws.title, coord, etiqueta, valor))
    return cel


def formula(ws, coord, txt, fmt=None, bold=None, destacar=False):
    """Fórmula registrada. Una celda que es SÓLO un SUMPRODUCT lleva «+0» al final: pycel
    devuelve 0 al agregar (SUM, MAX, SUMIF) celdas cuyo resultado es un SUMPRODUCT a secas
    (medido el 03-10-2026 en este libro: MAX y SUM de la rejilla daban 0 con la rejilla
    bien calculada; con «+0» dan lo correcto). En Excel el «+0» no cambia nada."""
    if txt.startswith('=SUMPRODUCT(') and txt.count('SUMPRODUCT') == 1 and txt.endswith(')'):
        txt += '+0'
    cel = motor.f(ws, coord, txt, fmt=fmt, bold=bold)
    if destacar:
        crema(cel)
    return cel


def nota_legal(ws, coord, pid):
    txt = D.nota_legal(pid)
    if not txt:
        raise SystemExit('nota_legal(%r) vacía: gate_legal() debería haber abortado' % pid)
    ws[coord].comment = Comment(txt, 'AI Chef Pro', height=130, width=470)
    NOTAS_LEGALES.append((ws.title, coord, pid))
    return txt


RX_ID = D.RX_ID


def nota_fuente(ws, coord, fuente, extra=None):
    """Comentario de procedencia de un dato de SECTOR (CUS-*): id, fiabilidad y URL."""
    partes = []
    for pid in RX_ID.findall(fuente or ''):
        if pid in D.IDS_PROHIBIDOS:
            raise SystemExit('El id %s está PROHIBIDO (SPEC §5.1) y no puede ser fuente' % pid)
        if pid.startswith('CUN-') or pid.startswith('CHN-'):
            continue
        f = D.ficha_comun(pid)
        if not f:
            raise SystemExit('El id %s no existe en el JSON común' % pid)
        url = str(f.get('url') or '').strip()
        fiab = str(f.get('fiabilidad') or '').strip()
        partes.append('%s%s%s' % (pid, ' · fiabilidad ' + fiab if fiab else '',
                                  ' · ' + url if url else ''))
    if extra:
        partes.append(extra)
    if not partes:
        return None
    ws[coord].comment = Comment('Fuente: ' + ' | '.join(partes), 'AI Chef Pro',
                                height=110, width=440)
    return partes


def texto_fuente(fuente):
    if not fuente:
        return ''
    if fuente.startswith('supuesto'):
        return 'supuesto declarado de El Molinete' + fuente[len('supuesto'):]
    if fuente.startswith('heredado'):
        return 'heredado de la hermana o del motor (' + fuente[len('heredado: '):] + ')'
    return fuente


def setup(ws, apaisado=True, titulos=None):
    CC.pagina(ws, apaisado=apaisado, titulos=titulos)


def version_al_pie(ws, fila, col_='A'):
    motor.val(ws, '%s%d' % (col_, fila), CC.VERSION_LINE)
    ws['%s%d' % (col_, fila)].font = Font(size=8, italic=True)


def semaforo(ws, coord, verdes_txt, rojos_txt, ambar_txt=()):
    voc = [(t, motor.CF_VERDE_BG, motor.CF_VERDE_FG) for t in verdes_txt]
    voc += [(t, motor.CF_AMBAR_BG, motor.CF_AMBAR_FG) for t in ambar_txt]
    voc += [(t, motor.CF_ROJO_BG, motor.CF_ROJO_FG) for t in rojos_txt]
    motor.semaforo_texto(ws, coord, tuple(voc))


# ==========================================================================
# Datos que el libro necesita (todos de `datos_ejemplo`)
# ==========================================================================
CRUCES_IN = [c for c in D.CRUCES if c['receptor'] == LIBRO]
CRUCES_OUT = [c for c in D.CRUCES if c['origen'] == LIBRO]
assert [c['x'] for c in CRUCES_IN] == ['X6', 'X7', 'X8'], [c['x'] for c in CRUCES_IN]
assert [c['calcula'] for c in CRUCES_OUT] == ['coste_plantilla_anual']
assert (CRUCES_OUT[0]['hoja_origen'], CRUCES_OUT[0]['celda_origen']) == (H_COS, 'L5')

PLANTILLA = D.PLANTILLA
TURNOS = list(D.TURNOS)
MINIMOS = D.MINIMOS_COBERTURA
DIAS = D.DIAS
N_P = len(PLANTILLA)
assert N_P == 5 and len(TURNOS) == len(D.TURNOS) and len(MINIMOS) == 6

#: Qué comprobación cuadra cada celda cruzada.
IDENTIDAD = {'_personas_pico': 'E', 'horas_refuerzo_pico': 'G', 'coste_hora_mano_obra': 'H'}
ETIQUETA_CRUCE = {
    '_personas_pico': ('Personas en la línea en el pico del fin de semana', 'personas', ENT),
    'horas_refuerzo_pico': ('Horas de refuerzo de Navidad y Reyes al año', 'h/año', DEC1),
    'coste_hora_mano_obra': ('Coste hora de mano de obra que usaste en el escandallo', 'EUR/h', EUR),
}
NOTA_CRUCE = {
    '_personas_pico': ('Sale de la capacidad de la línea del libro 1. Aquí manda el mínimo de '
                       'personas en la línea del pico en «Cuadrante por Franja»: si tu cuadrante '
                       'no las pone, aparecen huecos.'),
    'horas_refuerzo_pico': ('Días punta de Navidad y Reyes con una persona más por la mañana. '
                            'Se pagan al coste hora de la plantilla en «Coste de Plantilla».'),
    'coste_hora_mano_obra': ('El libro 3 va antes y lo supone para valorar la mano de obra de '
                             'cada ración. Aquí se cuadra con el que sale de tu plantilla.'),
}

# --- Parámetros: filas ------------------------------------------------------
P_CAB = 5
P_SEC_X = 6
CRUCE_ROW = dict((c['calcula'], P_SEC_X + 1 + i) for i, c in enumerate(CRUCES_IN))
PR = {}          # clave -> fila de «Parámetros» (se rellena al construir la hoja)
NIV_ROWS = {}    # nivel -> fila


def PB(clave):
    return ABS_(H_PAR, 'B%d' % PR[clave])


def XB(clave):
    """Celda verde del cruce `clave` (nombre de la función de `datos_ejemplo`)."""
    return ABS_(H_PAR, 'B%d' % CRUCE_ROW[clave])


TOLERANCIA = 0.05            # de los CONTRASTES (el coste hora de tu plantilla frente al del libro 3)
TOLERANCIA_CUADRE = 0.02     # la de la familia: filas de CUADRE contra lo que publica el origen
DIAS_ANIO = 365
DIAS_SEMANA = len(DIAS)
HORAS_DIA = 24.0
PRIMERA_MEDIA_HORA = 4.0
TRAMO_H = 0.5
N_SLOTS = 42
PAGAS_OFERTA = 12

# --- Cuadrante por Franja -----------------------------------------------------
# R-15: las horas por franja del libro 5 ya no se copian aquí (sólo alimentaban una suma de
# control). El resumen del cuadrante abre la hoja.
CR_SEC = 4
(CR_HAP, CR_HEX, CR_HCU, CR_HUE, CR_VER, CR_MAXL) = range(CR_SEC + 1, CR_SEC + 7)
CM_SEC = CR_MAXL + 2
CM_CAB = CM_SEC + 1
CM_INI = CM_CAB + 1
CM_FIN = CM_INI + len(MINIMOS) - 1
CM_PICO = CM_INI + 1                       # «SD 6-10»: el pico (X6)
assert MINIMOS[1][0] == 'SD' and MINIMOS[1][4] == D.dato(D.PRODUCCION, 'personas_en_linea_pico')
CT_SEC = CM_FIN + 2
CT_CAB = CT_SEC + 1
CT_INI = CT_CAB + 1
CT_FIN = CT_INI + len(TURNOS) - 1
CG_SEC = CT_FIN + 3
CG_GRUPO = CG_SEC + 1
CG_CAB = CG_GRUPO + 1
CG_INI = CG_CAB + 1
CG_FIN = CG_INI + N_SLOTS - 1
CG_TOT = CG_FIN + 1
G_PRES, G_LIN, G_MIN, G_HUE = 'C', 'K', 'S', 'AA'   # primera columna de cada rejilla

# --- Nocturnidad ----------------------------------------------------------------
N_GRUPO = 6
N_CAB = 7
N_INI = 8
N_FIN = N_INI + N_P - 1
N_NOTAS = N_FIN + 2
N_M1_CAB = N_NOTAS + 6
N_M1 = N_M1_CAB + 1
N_M2_CAB = N_M1 + N_P + 1
N_M2 = N_M2_CAB + 1
N_M3_CAB = N_M2 + N_P + 1
N_M3 = N_M3_CAB + 1

# --- Coste de Plantilla -----------------------------------------------------------
K_CAB = 8
K_INI = 9
K_FIN = K_INI + N_P - 1
K_TOT = K_FIN + 1
K_SEC = K_TOT + 2
(K_PUESTOS, K_HORAS, K_CHORA, K_HEMP, K_SEMS, K_HSUS, K_HREF, K_CEXT, K_TOTAL) = range(K_SEC + 1, K_SEC + 10)
K_MER_SEC = K_TOTAL + 3
K_MER_CAB = K_MER_SEC + 2
K_MER_INI = K_MER_CAB + 1
K_MER_FIN = K_MER_INI + len(D.SALARIOS_MERCADO) - 1

# --- Horas del Titular -----------------------------------------------------------
(T_HSEM, T_JOR, T_EXTRA, T_DIAS, T_MADRU, T_SEMS, T_HANIO, T_EXANIO, T_VALOR, T_VERED,
 T_DEC) = range(6, 17)


# ==========================================================================
# Hoja «Instrucciones»
# ==========================================================================
def _orden_relleno():
    partes = ['%d (%s)' % (n, D.LIBROS[n]) for n in D.ORDEN_RELLENO if n != 7]
    return ('Orden de relleno del paquete: ' + ', luego '.join(partes)
            + '. El libro 7 (%s) no depende de ninguno: rellénalo cuando quieras. '
              'Este es el libro 8: rellénalo DESPUÉS del 3, el 1 y el 5, que le pasan sus '
              'cifras, y ANTES del 6, que copia el coste anual de plantilla.' % D.LIBROS[7])


def _cruces_texto():
    lineas = []
    for c in CRUCES_IN:
        lineas.append('Parámetros!B%d - %s. Viene del libro %d: %s!%s!%s.' % (
            CRUCE_ROW[c['calcula']], c['concepto'], c['origen'], c['fichero_origen'],
            c['hoja_origen'], c['celda_origen']))
    for c in CRUCES_OUT:
        lineas.append('%s!%s - %s. Lo copia: libro %d (%s, %s).' % (
            c['hoja_origen'], c['celda_origen'], c['concepto'], c['receptor'],
            c['fichero_receptor'], c['x']))
    return lineas


PASOS = [
    '1. Hoja «Parámetros». Arriba, tres celdas verdes que copias a mano de los libros 1, 5 '
    'y 3 (cada una dice de qué celda exacta). Debajo, tu convenio (provincia, tablas, pagas, '
    'plus de convenio y salario base por nivel), las dos figuras de nocturnidad con sus '
    'franjas, el SMI y la Seguridad Social. Al pie, primero el CONTRASTE de tus copias con '
    'tu plantilla y después el CUADRE: anota lo que publica hoy la celda de origen de cada '
    'cifra y, si lo que usas se separa, su fila se pone en rojo.',
    '2. Hoja «Cuadrante por Franja». Los mínimos de cobertura (cuántas personas hacen falta '
    'en cada tramo y cuántas en la línea de fritura) y los turnos de la SEMANA TIPO, persona '
    'a persona. La rejilla de medias horas cuenta quién está y marca cada hueco: tiene que '
    'dar cero.',
    '3. Hoja «Nocturnidad ET y Convenio». Para cada persona, las dos figuras por separado: '
    'si es trabajador nocturno según el Estatuto (sí, no o revísalo con tu asesor) y '
    'cuántas horas cobra con el plus del convenio y cuánto cuesta ese plus al año.',
    '4. Hoja «Coste de Plantilla». Jornada, nivel y si cotiza la Seguridad Social de empresa '
    'de cada puesto; su salario, el plus, el suelo del SMI y el coste anual. Abajo, el coste '
    'hora, las sustituciones de vacaciones, el refuerzo de Navidad y el COSTE ANUAL DE '
    'PLANTILLA que copia el libro 6.',
    '5. Hoja «Identifica tu Convenio». El método para encontrar el tuyo, con Madrid como '
    'ejemplo. Rellena tu respuesta en cada paso.',
    '6. Hoja «Horas del Titular». Las horas que haces tú, las que van por encima de una '
    'jornada completa, la madrugada que asumes y lo que valen al coste hora.',
]

NOTAS_LIBRO = [
    'DOS FIGURAS DISTINTAS QUE NO SE MEZCLAN. El Estatuto de los Trabajadores define al '
    'TRABAJADOR NOCTURNO por las horas entre las 22:00 y las 6:00 (al menos tres de su '
    'jornada diaria, o un tercio de la anual): sólo a él le alcanzan la prohibición de horas '
    'extra y el tope de ocho horas de media. El PLUS DE NOCTURNIDAD lo paga el convenio por '
    'las horas de sus franjas (en el de Madrid, de 22:00 a 0:00 y de 0:00 a 8:00). Quien hace '
    'el turno de las 4:30 cobra plus y, con hora y media en el periodo del Estatuto, no es '
    'trabajador nocturno.',
    'CONVENIO DE EJEMPLO, NO EL TUYO. Las cuantías son las del convenio de hostelería de la '
    'Comunidad de Madrid, clase C, tablas de 2025: el convenio venció el 31-12-2025 y está en '
    'negociación; se aplican sus tablas mientras no se publique el nuevo. Los niveles de cada '
    'puesto son una asimilación de funciones propuesta: compruébala con tu asesor. La hoja '
    '«Identifica tu Convenio» te enseña a buscar el tuyo.',
    'EL SUELO ES EL SMI. El coste de cada puesto se calcula sobre el mayor de dos importes: '
    'su salario (base, pagas, plus de convenio y plus de nocturnidad, a su jornada) o el SMI '
    'anual a esa misma jornada. Si el salario queda por debajo, el semáforo lo dice y el '
    'coste usa el SMI.',
    'SEMANA TIPO, NO EL CUADRANTE DE CADA SEMANA. Este libro calcula la semana modelo que '
    'usa el plan: horas por franja, las dos figuras de nocturnidad, el plus y el coste anual '
    'por puesto. No reproduce el cuadrante semana a semana por empleado, con cambios, '
    'vacaciones y fichajes, ni las nóminas: para eso está el Kit Gestión de Personal y '
    'Turnos (aichef.pro/kit-gestion-personal).',
    'EL REGISTRO DE JORNADA ES DIARIO. Cada día, hora de inicio y de fin de cada persona, y '
    'se guarda cuatro años. Los turnos de este libro son el plan; el registro es lo que pasó.',
    'LOS SUPUESTOS NO SON DATOS DEL SECTOR. La plantilla, los turnos, los mínimos de '
    'cobertura, las 14 pagas y la tolerancia de los cuadres son supuestos declarados de El '
    'Molinete. Los salarios de mercado de la hoja de coste son ofertas reales con las pagas '
    'sin declarar: sirven para saber si estás en mercado; lo que te obliga es tu convenio.',
]

CADENCIA = ('Cada cuánto se usa este libro: al montar la plantilla, cada vez que cambies un '
            'turno o contrates, y cuando se publiquen las tablas nuevas de tu convenio o el '
            'SMI del año. Si cambia el coste anual de plantilla del bloque «Lo que Copian '
            'Otros Libros» de «Coste de Plantilla», cópialo de nuevo en el libro 6.')


def hoja_instrucciones(wb):
    ws = wb.create_sheet(H_INS, 0)
    ws.column_dimensions['A'].width = 100.0
    cabecera_hoja(ws, TITULO)
    motor.val(ws, 'A3', 'Para qué sirve: saber cuánta gente necesitas, si alguien es '
                        'trabajador nocturno, cuánto cuesta el plus de la madrugada, cuánto '
                        'cuesta tu plantilla al año y cuántas horas haces tú.')
    ws['A3'].font = Font(italic=True, size=9)
    fila = 5
    seccion(ws, 'A%d' % fila, 'Instrucciones de Uso')
    fila += 1
    entrada(ws, 'A%d' % fila, motor.NOTA_VERDES, etiqueta='Leyenda de celdas verdes')
    fila += 1
    motor.val(ws, 'A%d' % fila, 'Las celdas en crema son resultados que el libro calcula; las '
                                'blancas con fórmula, también. No escribas encima.', wrap=True)
    fila += 2
    for paso in PASOS:
        motor.val(ws, 'A%d' % fila, paso, wrap=True)
        ws.row_dimensions[fila].height = 58
        fila += 1
    fila += 1
    seccion(ws, 'A%d' % fila, 'Orden de Relleno y Cifras que Viajan')
    fila += 1
    motor.val(ws, 'A%d' % fila, _orden_relleno(), wrap=True)
    ws.row_dimensions[fila].height = 58
    fila += 1
    motor.val(ws, 'A%d' % fila,
              'Ningún libro se vincula con otro por fórmula. Lo que viene de otro libro tiene en '
              '«Parámetros» una celda verde para copiarlo a mano, con la celda exacta de origen, '
              'y una fila de CUADRE que avisa si se ha quedado vieja frente a tu plantilla. Lo '
              'que este libro pasa a otro está en el bloque «Lo que Copian Otros Libros» (rótulo '
              'en la columna K, valor en la L).', wrap=True)
    ws.row_dimensions[fila].height = 58
    fila += 1
    for linea in _cruces_texto():
        motor.val(ws, 'A%d' % fila, linea, wrap=True)
        ws['A%d' % fila].font = Font(size=9)
        ws.row_dimensions[fila].height = 28
        fila += 1
    fila += 1
    seccion(ws, 'A%d' % fila, 'Lo que Conviene Saber Antes de Empezar')
    fila += 1
    for txt in NOTAS_LIBRO:
        motor.val(ws, 'A%d' % fila, txt, wrap=True)
        ws.row_dimensions[fila].height = 72
        if txt.startswith('DOS FIGURAS'):
            nota_legal(ws, 'A%d' % fila, 'CUN-29')
        if txt.startswith('CONVENIO DE EJEMPLO'):
            nota_legal(ws, 'A%d' % fila, 'CUN-V01')
        if txt.startswith('EL SUELO'):
            nota_legal(ws, 'A%d' % fila, 'CHN-67')
        if txt.startswith('EL REGISTRO'):
            nota_legal(ws, 'A%d' % fila, 'CHN-68')
        fila += 1
    fila += 1
    motor.val(ws, 'A%d' % fila, CADENCIA, wrap=True)
    ws.row_dimensions[fila].height = 58
    fila += 2
    motor.val(ws, 'A%d' % fila, CC.NOTA_DESPROTEGER, wrap=True)
    fila += 2
    motor.val(ws, 'A%d' % fila, CC.BIO, wrap=True)
    fila += 1
    version_al_pie(ws, fila)
    setup(ws, apaisado=False)
    return ws


# ==========================================================================
# Hoja «Parámetros»
# ==========================================================================
def fila_param(ws, fila, etiqueta, valor, fmt, unidad, origen, nota, pct=False, form=None,
               minimo=0, maximo=None, legal=None, fuente_sector=None, texto_libre=False):
    motor.val(ws, 'A%d' % fila, etiqueta, wrap=True)
    if form is not None:
        cel = formula(ws, 'B%d' % fila, form, fmt=fmt, bold=True, destacar=True)
    else:
        cel = entrada(ws, 'B%d' % fila, valor, fmt=fmt, etiqueta=etiqueta)
        if pct:
            motor.dv_porcentaje(ws, [cel.coordinate])
        elif not texto_libre and isinstance(valor, (int, float)):
            motor.dv_numerica(ws, [cel.coordinate], minimo=minimo, maximo=maximo)
    motor.val(ws, 'C%d' % fila, unidad)
    motor.val(ws, 'D%d' % fila, origen, wrap=True)
    ws['D%d' % fila].font = Font(size=8, color=GRIS)
    gris(ws, 'E%d' % fila, nota)
    ws.row_dimensions[fila].height = 44
    if legal:
        nota_legal(ws, cel.coordinate, legal)
    elif fuente_sector:
        nota_fuente(ws, cel.coordinate, fuente_sector)
    return cel


def _veredicto(cond_ko, guardas, ko):
    return '=IF(AND(%s),IF(%s,"%s","COHERENTE"),"")' % (
        ','.join('ISNUMBER(%s)' % g for g in guardas), cond_ko, ko)


KO = {
    'E': ('INCOHERENTE: tu cuadrante no pone nunca en la línea las personas que el libro 1 dice '
          'que hacen falta en el pico; revisa los turnos o vuelve a copiar el libro 1'),
    'G': ('INCOHERENTE: las horas de refuerzo no caben en las horas de apertura del año; vuelve '
          'a copiar el libro 5'),
    'H': ('INCOHERENTE: el coste hora que usaste en el escandallo del libro 3 se separa del que '
          'sale de tu plantilla; copia este en el libro 3 y vuelve a copiarlo aquí'),
}


def _p(ws, fila, clave, etiqueta, valor, fmt, unidad, origen, nota, **kw):
    PR[clave] = fila
    return fila_param(ws, fila, etiqueta, valor, fmt, unidad, origen, nota, **kw)


def hoja_parametros(wb):
    ws = wb.create_sheet(H_PAR)
    cabecera_hoja(ws, H_PAR, 'Todo lo que el resto de hojas da por sabido vive aquí. Ninguna '
                             'fórmula de este libro lleva dentro un salario, un porcentaje, '
                             'una hora ni un umbral.')
    for letra, ancho in (('A', 54), ('B', 18), ('C', 13), ('D', 46), ('E', 66)):
        ws.column_dimensions[letra].width = ancho
    encabezados(ws, P_CAB, [('A', 'Parámetro', None), ('B', 'Valor', None),
                            ('C', 'Unidad', None), ('D', 'De dónde sale', None),
                            ('E', 'Nota', None)], alto=24)

    seccion(ws, 'A%d' % P_SEC_X, 'Lo que Copias de los Libros 1, 5 y 3')
    for c in CRUCES_IN:
        clave = c['calcula']
        etq, uni, fmt = ETIQUETA_CRUCE[clave]
        valor = c['valor_defecto']
        if fmt == ENT:
            valor = int(round(valor))
        cel = fila_param(ws, CRUCE_ROW[clave], etq, valor, fmt, uni, D.rotulo_cruce(c),
                         NOTA_CRUCE[clave], minimo=0)
        CRUCE_VERDES.append((ws.title, cel.coordinate, c['x'], c['concepto']))

    C = D.CONVENIO
    fila = CRUCE_ROW['coste_hora_mano_obra'] + 2
    seccion(ws, 'A%d' % fila, 'Tu Convenio (Ejemplo: Madrid, Clase C, Tablas 2025)')
    fila += 1
    v, f, n = C['provincia']
    _p(ws, fila, 'provincia', 'Provincia de tu local', v, None, '', texto_fuente(f),
       'Pon la tuya: de ella depende el convenio que te toca.')
    fila += 1
    v, f, n = C['codigo']
    _p(ws, fila, 'codigo', 'Código REGCON de tu convenio', v, '@', '', f,
       'El del ejemplo es el de hostelería de la Comunidad de Madrid. Busca el tuyo con la '
       'hoja «Identifica tu Convenio».', legal='CUN-39')
    fila += 1
    v, f, n = C['denominacion']
    _p(ws, fila, 'denominacion', 'Nombre de tu convenio', v, None, '', f, n)
    fila += 1
    v, f, n = C['clase']
    _p(ws, fila, 'clase', 'Clase o grupo en el que entra tu local', v, None, '', f, n)
    fila += 1
    v, f, n = C['tablas']
    _p(ws, fila, 'tablas', 'Tablas salariales que usas', v, None, '', f, n, legal='CUN-V01')
    fila += 1
    v, f, n = C['pagas']
    _p(ws, fila, 'pagas', 'Pagas al año', v, ENT, 'pagas', texto_fuente(f), n, minimo=1, maximo=16)
    fila += 1
    v, f, n = C['plus_convenio_mes']
    _p(ws, fila, 'plus_conv', 'Plus de convenio', v, EUR, 'EUR/mes', f, n, legal='CUN-39')
    fila += 1
    v, f, n = C['meses_plus_convenio']
    _p(ws, fila, 'meses_plus', 'Meses que se cobra el plus de convenio', v, ENT, 'meses', f, n,
       minimo=0, maximo=14)
    fila += 1
    v, f, n = C['smi_anual']
    _p(ws, fila, 'smi_anual', 'SMI anual (suelo de cada puesto a jornada completa)', v, EUR,
       'EUR/año', 'CHN-67', n, legal='CHN-67')
    fila += 1
    v, f, n = C['smi_mensual']
    _p(ws, fila, 'smi_mes', 'SMI mensual', v, EUR, 'EUR/mes', f,
       'Caduca el 31-12-2026. Sirve para leer una oferta de mercado mes a mes.', legal='CHN-67')
    fila += 1
    v, f, n = D.PARAMS['ss_empresa']
    _p(ws, fila, 'ss', 'Seguridad Social a cargo de la empresa', v, PCT, 'del bruto',
       texto_fuente(f), n, pct=True)
    fila += 1
    v, f, n = C['jornada_anual_h']
    _p(ws, fila, 'jor_anual', 'Jornada anual de un contrato a jornada completa', v, ENT, 'h/año',
       texto_fuente(f), n, minimo=1)
    fila += 1
    v, f, n = C['horas_semana_completa']
    _p(ws, fila, 'h_sem', 'Horas a la semana de una jornada completa (1,0)', v, DEC1, 'h/semana',
       texto_fuente(f), n, minimo=1)

    fila += 2
    seccion(ws, 'A%d' % fila, 'Salario Base Mensual por Nivel (Tabla de Ejemplo)')
    fila += 1
    gris(ws, 'A%d' % fila, 'En la columna A, el nombre del nivel tal como lo eliges en «Coste '
                           'de Plantilla»; en la B, su salario base de 2025. Si tu convenio '
                           'tiene otros niveles, cambia el nombre y el importe.', alto=30)
    for nivel, (base, fuente) in sorted(D.NIVELES.items(), key=lambda kv: -kv[1][0]):
        fila += 1
        NIV_ROWS[nivel] = fila
        motor.val(ws, 'A%d' % fila, nivel, bold=True)
        cel = entrada(ws, 'B%d' % fila, base, fmt=EUR, etiqueta='Salario base nivel ' + nivel)
        motor.dv_numerica(ws, [cel.coordinate], minimo=0)
        motor.val(ws, 'C%d' % fila, 'EUR/mes')
        motor.val(ws, 'D%d' % fila, fuente + ' · ' + D.dato(C, 'tablas'), wrap=True)
        ws['D%d' % fila].font = Font(size=8, color=GRIS)
        gris(ws, 'E%d' % fila, 'Nivel %s de la clase C del ejemplo. El puesto que encajas aquí '
                               'es una %s.' % (nivel, D.ASIMILACION))
        ws.row_dimensions[fila].height = 32
        nota_legal(ws, cel.coordinate, 'CUN-39')

    ET = D.TRABAJO_NOCTURNO_ET
    fila += 2
    seccion(ws, 'A%d' % fila, 'Figura 1: Trabajador Nocturno del Estatuto (Art. 36.1)')
    fila += 1
    v, f, n = ET['inicio']
    _p(ws, fila, 'et_ini', 'Empieza el periodo nocturno del Estatuto a las', v, DEC1, 'h', f, n,
       legal='CUN-29', maximo=24)
    fila += 1
    v, f, n = ET['fin']
    _p(ws, fila, 'et_fin', 'Termina el periodo nocturno del Estatuto a las (hora del día '
                          'siguiente)', v - HORAS_DIA, DEC1, 'h', f,
       'Las 6:00 de la mañana siguiente. Escribe la hora del reloj: 6.', legal='CUN-29', maximo=24)
    fila += 1
    v, f, n = ET['horas_diarias_minimas']
    _p(ws, fila, 'et_min', 'Horas diarias en ese periodo para ser trabajador nocturno', v, DEC1,
       'h/día', f, n, legal='CUN-29')
    fila += 1
    v, f, n = ET['fraccion_jornada_anual']
    _p(ws, fila, 'et_tercio', 'O parte de la jornada anual en ese periodo', v, PCT, 'de la jornada',
       f, n, pct=True, legal='CUN-29')

    fila += 2
    seccion(ws, 'A%d' % fila, 'Figura 2: Plus de Nocturnidad del Convenio (Franjas en Verde)')
    fila += 1
    _p(ws, fila, 'p1_ini', 'Franja 1 del plus: desde las', 22.0, DEC1, 'h', 'CUN-40',
       'Primera franja del convenio de ejemplo: de 22:00 a 0:00.', legal='CUN-40', maximo=24)
    fila += 1
    _p(ws, fila, 'p1_fin', 'Franja 1 del plus: hasta las (24 = medianoche)', 24.0, DEC1, 'h',
       'CUN-40', 'Medianoche se escribe 24.', legal='CUN-40', maximo=24)
    fila += 1
    v, f, n = C['plus_noche_22_0']
    _p(ws, fila, 'p1_pct', 'Franja 1 del plus: recargo sobre el salario base por hora', v, PCT,
       'por hora', f, n, pct=True, legal='CUN-40')
    fila += 1
    _p(ws, fila, 'p2_ini', 'Franja 2 del plus: desde las (0 = medianoche)', 0.0, DEC1, 'h',
       'CUN-40', 'Segunda franja: de 0:00 a 8:00.', legal='CUN-40', maximo=24)
    fila += 1
    _p(ws, fila, 'p2_fin', 'Franja 2 del plus: hasta las', 8.0, DEC1, 'h', 'CUN-40',
       'Las horas de esta franja son también las de la madrugada del titular.', legal='CUN-40',
       maximo=24)
    fila += 1
    v, f, n = C['plus_noche_0_8']
    _p(ws, fila, 'p2_pct', 'Franja 2 del plus: recargo sobre el salario base por hora', v, PCT,
       'por hora', f, n, pct=True, legal='CUN-40')
    fila += 1
    v, f, n = C['umbral_jornada_nocturna_h']
    _p(ws, fila, 'umbral', 'Horas entre el inicio de la franja 1 y el final de la 2 para que la '
                          'jornada cuente entera como nocturna', v, DEC1, 'h/día', f, n,
       legal='CUN-40')

    fila += 2
    seccion(ws, 'A%d' % fila, 'Calendario, Sustituciones y Cuadrante')
    fila += 1
    _p(ws, fila, 'dia', 'Horas de un día', HORAS_DIA, DEC1, 'h', 'calendario',
       'Los turnos que pasan de medianoche se escriben sumando un día: la 1:00 es 25.', minimo=1)
    fila += 1
    _p(ws, fila, 'anio', 'Días naturales del año', DIAS_ANIO, ENT, 'días', 'calendario', '',
       minimo=1)
    fila += 1
    _p(ws, fila, 'dsem', 'Días de la semana', DIAS_SEMANA, ENT, 'días', 'calendario', '', minimo=1)
    fila += 1
    v, f, n = D.NEGOCIO['dias_cierre_anio']
    _p(ws, fila, 'cierre', 'Días al año con el local cerrado', v, ENT, 'días', texto_fuente(f), n)
    fila += 1
    _p(ws, fila, 'sem_ab', 'Semanas que el local abre al año', None, DEC, 'semanas', 'calculado',
       'Días abiertos entre los días de la semana.',
       form='=IFERROR((%s-%s)/%s,"")' % (PB('anio'), PB('cierre'), PB('dsem')))
    fila += 1
    _p(ws, fila, 'sem_tr', 'Semanas que trabaja un contrato al año', None, DEC, 'semanas',
       'calculado', 'Jornada anual entre las horas de una semana completa: el resto son '
       'vacaciones y descansos por festivo.',
       form='=IFERROR(%s/%s,"")' % (PB('jor_anual'), PB('h_sem')))
    fila += 1
    v, f, n = D.PARAMS['cubrir_vacaciones_con_sustitutos']
    _p(ws, fila, 'cubrir', '¿Cubres vacaciones y festivos de la plantilla con sustitutos?',
       SI if v else NO, None, 'sí / no', texto_fuente(f), n)
    fila += 1
    _p(ws, fila, 'pagas_oferta', 'Pagas con las que lees una oferta de mercado que no las dice',
       PAGAS_OFERTA, ENT, 'pagas', 'supuesto declarado de este libro',
       'La lectura prudente de una oferta sin pagas: doce. Con doce, 1.300 euros al mes quedan '
       'por debajo del SMI anual.', minimo=1, maximo=16)
    fila += 1
    _p(ws, fila, 'slot0', 'Primera media hora de la rejilla del cuadrante', PRIMERA_MEDIA_HORA,
       DEC1, 'h', 'supuesto declarado de este libro',
       'La rejilla de «Cuadrante por Franja» cubre %d tramos desde aquí. Ponla antes del primer '
       'turno de la mañana.' % N_SLOTS, maximo=24)
    fila += 1
    _p(ws, fila, 'tramo', 'Duración de cada tramo de la rejilla', TRAMO_H, DEC, 'h',
       'supuesto declarado de este libro', 'Media hora: con turnos que empiezan a y media.',
       minimo=0.25, maximo=1)
    fila += 1
    _p(ws, fila, 'tol', 'Tolerancia de los contrastes', TOLERANCIA, PCT, '',
       'supuesto declarado de este libro',
       'Cuánto pueden separarse dos cifras que deberían coincidir (el coste hora de tu plantilla '
       'frente al del escandallo) antes de que la hoja avise.', pct=True)
    fila += 1
    _p(ws, fila, 'tolc', 'Tolerancia de las filas de CUADRE', TOLERANCIA_CUADRE, PCT, '',
       'heredado de la Guía de la Chocolatería (tolerancia de las filas de CUADRE)',
       'Por encima de esta desviación, la fila de CUADRE avisa de que lo que has copiado no es lo '
       'que publica el libro de origen.', pct=True)

    # --- CONTRASTE entre tus copias y tu plantilla -------------------------------------
    fila += 2
    seccion(ws, 'A%d' % fila, 'Contraste entre tus Copias y tu Plantilla')
    t = PB('tol')
    comprob = (
        ('c_maxl', 'Máximo de personas en la línea a la vez en tu cuadrante',
         '=%s' % ABS_(H_CUA, 'B%d' % CR_MAXL), ENT, 'personas'),
        ('c_minimos', 'Horas de apertura con mínimo de cobertura en tu cuadrante',
         '=%s' % ABS_(H_CUA, 'B%d' % CR_HAP), DEC1, 'h/semana'),
        ('c_anio', 'Horas de apertura al año (las de la semana por las semanas abiertas)', None, DEC1,
         'h/año'),
        ('c_chora', 'Coste hora que sale de tu plantilla («Coste de Plantilla»)',
         '=%s' % ABS_(H_COS, 'B%d' % K_CHORA), EUR, 'EUR/h'),
        ('c_sep_h', 'Separación entre el coste hora del libro 3 y el de tu plantilla', None, PCT, ''),
    )
    for clave, etq, form, fmt, uni in comprob:
        fila += 1
        PR[clave] = fila
        if clave == 'c_anio':
            form = '=IFERROR(B%d*%s,"")' % (PR['c_minimos'], PB('sem_ab'))
        elif clave == 'c_sep_h':
            form = '=IFERROR(ABS(%s/B%d-1),"")' % (XB('coste_hora_mano_obra'), PR['c_chora'])
        motor.val(ws, 'A%d' % fila, etq, wrap=True)
        ws['A%d' % fila].font = Font(size=9)
        formula(ws, 'B%d' % fila, form, fmt=fmt)
        motor.val(ws, 'C%d' % fila, uni)
        ws.row_dimensions[fila].height = 30
    gris(ws, 'D%d' % PR['c_chora'], 'El libro 3 lo supone antes de que exista la plantilla; '
                                    'aquí se comprueba.')

    fila += 2
    encabezados(ws, fila, [('A', 'Contraste', None), ('B', 'Veredicto', None),
                           ('C', 'Celda que compara', None), ('D', 'Qué compara', None),
                           ('E', 'Si no es coherente', None)], alto=24, congelar=False)
    que_compara = {
        'E': 'Personas en la línea del libro 1 frente al máximo que pone tu cuadrante (fila %d)'
             % PR['c_maxl'],
        'G': 'Horas de refuerzo frente a las horas de apertura del año (fila %d)' % PR['c_anio'],
        'H': 'Coste hora del libro 3 frente al de tu plantilla (filas %d y %d)'
             % (PR['c_chora'], PR['c_sep_h']),
    }
    corto = {'E': 'personas en la línea del pico', 'G': 'horas de refuerzo',
             'H': 'coste hora del escandallo'}
    for clave in ('_personas_pico', 'horas_refuerzo_pico', 'coste_hora_mano_obra'):
        fila += 1
        ident = IDENTIDAD[clave]
        PR['contraste_' + clave] = fila
        motor.val(ws, 'A%d' % fila, 'Contraste %s · %s' % (ident, corto[ident]), wrap=True)
        ws['A%d' % fila].font = Font(bold=True, size=9)
        if ident == 'E':
            form = _veredicto('B%d<%s' % (PR['c_maxl'], XB(clave)),
                              ['B%d' % PR['c_maxl'], XB(clave)], KO['E'])
        elif ident == 'G':
            form = _veredicto('OR(%s<0,%s>B%d)' % (XB(clave), XB(clave), PR['c_anio']),
                              [XB(clave), 'B%d' % PR['c_anio']], KO['G'])
        else:
            form = _veredicto('B%d>%s' % (PR['c_sep_h'], t), ['B%d' % PR['c_sep_h']], KO['H'])
        cel = formula(ws, 'B%d' % fila, form, bold=True)
        cel.alignment = Alignment(wrap_text=True, vertical='top')
        semaforo(ws, 'B%d' % fila, ('COHERENTE',), (KO[ident],))
        motor.val(ws, 'C%d' % fila, 'B%d' % CRUCE_ROW[clave])
        gris(ws, 'D%d' % fila, que_compara[ident])
        gris(ws, 'E%d' % fila, 'Revisa tu cuadrante, tus horas de refuerzo o vuelve a copiar la '
                               'cifra de su libro de origen.')
        ws.row_dimensions[fila].height = 48
        CONTRASTES.append((ws.title, 'B%d' % fila))

    # --- CUADRE: una fila por celda cruzada, contra lo que publica el origen (R-01) -------
    fila += 2
    seccion(ws, 'A%d' % fila, 'Cuadre de lo que Has Copiado')
    fila += 1
    encabezados(ws, fila, [('A', 'Cuadre', None),
                           ('B', 'Lo que publica hoy la celda de origen (anótalo)', None),
                           ('C', 'Desviación', None), ('D', 'Veredicto', None),
                           ('E', 'Si no cuadra', None)], alto=36, congelar=False)
    tc = PB('tolc')
    for c in CRUCES_IN:
        fila += 1
        clave = c['calcula']
        PR['cuadre_' + clave] = fila
        etq, _uni, fmt = ETIQUETA_CRUCE[clave]
        motor.val(ws, 'A%d' % fila, 'CUADRE · %s (%s)' % (etq, c['x']), wrap=True)
        ws['A%d' % fila].font = Font(bold=True, size=9)
        valor = c['valor_defecto']
        if fmt == ENT:
            valor = int(round(valor))
        cel = entrada(ws, 'B%d' % fila, valor, fmt=fmt,
                      etiqueta='%s: lo que publica el libro %d' % (c['x'], c['origen']))
        motor.dv_numerica(ws, [cel.coordinate], minimo=0)
        formula(ws, 'C%d' % fila, '=IFERROR(ABS(%s-B%d)/B%d,"")' % (XB(clave), fila, fila), fmt=PCT)
        ko = 'REVISA: no es lo que publica el libro %d' % c['origen']
        cel = formula(ws, 'D%d' % fila, '=IF(NOT(ISNUMBER(C%d)),"",IF(C%d<=%s,"CUADRA","%s"))'
                      % (fila, fila, tc, ko), bold=True)
        cel.alignment = Alignment(wrap_text=True, vertical='top')
        semaforo(ws, 'D%d' % fila, ('CUADRA',), (ko,))
        gris(ws, 'E%d' % fila, 'Vuelve a copiar la cifra de su celda de origen (libro %d, %s!%s!%s) '
                               'en la fila %d, o anota aquí lo que publica hoy.'
             % (c['origen'], c['fichero_origen'], c['hoja_origen'], c['celda_origen'],
                CRUCE_ROW[clave]))
        ws.row_dimensions[fila].height = 40
        CRUCE_CUADRES.append((ws.title, 'D%d' % fila, c['x'], c['concepto']))

    fila += 3
    refs, fin = CC.bloque_listas(ws, fila, (('Sí o no', SI_NO),))
    CC.dv_rango(ws, ['B%d' % PR['cubrir']], refs['Sí o no'], 'Sustitutos',
                'Elige sí o no.')
    version_al_pie(ws, fin + 1)
    setup(ws, apaisado=True)
    return ws


# ==========================================================================
# Hoja «Cuadrante por Franja»
# ==========================================================================
def _ol(a, b, ini, fin):
    """Horas de solape entre el turno [a, b) y el tramo [ini, fin)."""
    return 'MAX(0,MIN(%s,%s)-MAX(%s,%s))' % (b, fin, a, ini)


def hoja_cuadrante(wb):
    ws = wb.create_sheet(H_CUA)
    cabecera_hoja(ws, H_CUA, 'La semana tipo del plan: mínimos de cobertura, turnos persona a '
                             'persona y una rejilla de medias horas que cuenta quién está y '
                             'marca los huecos. No es el cuadrante de cada semana.')
    for letra, ancho in (('A', 34), ('B', 13), ('C', 11), ('D', 11), ('E', 11), ('F', 11),
                         ('G', 11), ('H', 11), ('I', 11), ('J', 11), ('K', 7), ('L', 7)):
        ws.column_dimensions[letra].width = ancho

    # --- resumen -----------------------------------------------------------------------
    seccion(ws, 'A%d' % CR_SEC, 'Resumen del Cuadrante')
    rng_min = '%s%d:%s%d' % (G_MIN, CG_INI, col(G_MIN, 6), CG_FIN)
    rng_hue = '%s%d:%s%d' % (G_HUE, CG_INI, col(G_HUE, 6), CG_FIN)
    rng_lin = '%s%d:%s%d' % (G_LIN, CG_INI, col(G_LIN, 6), CG_FIN)
    resumen = (
        (CR_HAP, 'Horas de apertura con mínimo de cobertura (a la semana)',
         '=COUNTIF(%s,">0")*%s' % (rng_min, PB('tramo')), DEC1),
        (CR_HEX, 'Horas de trabajo que piden los mínimos (personas por horas)',
         '=SUM(%s)*%s' % (rng_min, PB('tramo')), DEC1),
        (CR_HCU, 'Horas de trabajo que pone tu cuadrante', '=SUM(F%d:F%d)' % (CT_INI, CT_FIN), DEC1),
        (CR_HUE, 'Medias horas sin cubrir (huecos de cobertura)', '=SUM(%s)' % rng_hue, ENT),
        (CR_VER, 'Veredicto de cobertura',
         '=IF(ISNUMBER(B%d),IF(B%d>0,"HAY HUECOS: mira la rejilla de huecos, abajo",'
         '"Sin huecos: el cuadrante cubre todos los mínimos"),"")' % (CR_HUE, CR_HUE), None),
        (CR_MAXL, 'Máximo de personas en la línea a la vez', '=MAX(%s)' % rng_lin, ENT),
    )
    for fila, etq, form, fmt in resumen:
        motor.val(ws, 'A%d' % fila, etq, wrap=True)
        formula(ws, 'B%d' % fila, form, fmt=fmt, bold=True, destacar=True)
        ws.row_dimensions[fila].height = 28
    ws.merge_cells('B%d:F%d' % (CR_VER, CR_VER))
    semaforo(ws, 'B%d' % CR_VER, ('Sin huecos: el cuadrante cubre todos los mínimos',),
             ('HAY HUECOS: mira la rejilla de huecos, abajo',))
    gris(ws, 'C%d' % CR_HCU, 'Más que lo que piden los mínimos: preparación de la masa antes de '
                             'abrir, solapes de relevo y el cierre.')

    # --- mínimos de cobertura ------------------------------------------------------------
    seccion(ws, 'A%d' % CM_SEC, 'Mínimos de Cobertura (1 = Ese Día se Exige)')
    cols = [('A', 'Tramo', None), ('B', 'Desde (h)', None), ('C', 'Hasta (h)', None),
            ('D', 'Personas', None), ('E', 'De ellas en la línea', None)]
    cols += [(col('F', j), d, None) for j, d in enumerate(DIAS)]
    encabezados(ws, CM_CAB, cols, alto=30, congelar=False)
    for i, (dias, ini, fin, n_tot, n_lin, nota_m) in enumerate(MINIMOS):
        fila = CM_INI + i
        motor.val(ws, 'A%d' % fila, nota_m, wrap=True)
        ws['A%d' % fila].font = Font(size=9)
        for letra, v in (('B', ini), ('C', fin), ('D', n_tot)):
            cel = entrada(ws, '%s%d' % (letra, fila), v, fmt=DEC1 if letra != 'D' else ENT,
                          etiqueta='mínimo ' + letra)
            motor.dv_numerica(ws, [cel.coordinate], minimo=0, maximo=48)
        if fila == CM_PICO:
            formula(ws, 'E%d' % fila, '=%s' % XB('_personas_pico'), fmt=ENT, bold=True)
            gris(ws, 'M%d' % fila, 'Personas en la línea del pico: las copia «Parámetros» del '
                                   'libro 1 (X6).')
        else:
            cel = entrada(ws, 'E%d' % fila, n_lin, fmt=ENT, etiqueta='mínimo en la línea')
            motor.dv_numerica(ws, [cel.coordinate], minimo=0, maximo=48)
        for j, d in enumerate(DIAS):
            cel = entrada(ws, '%s%d' % (col('F', j), fila), 1 if d in dias else 0, fmt='0',
                          etiqueta='día exigido')
            motor.dv_numerica(ws, [cel.coordinate], minimo=0, maximo=1,
                              mensaje='Escribe 1 si ese día se exige el mínimo y 0 si no.')
        ws.row_dimensions[fila].height = 30
    nota_fuente(ws, 'A%d' % CM_INI, '', extra='Supuestos declarados de El Molinete. '
                'Cada media hora de apertura tiene que caer '
                'en un tramo: si abres más horas, añade su tramo.')

    # --- turnos ----------------------------------------------------------------------
    seccion(ws, 'A%d' % CT_SEC, 'Turnos de la Semana Tipo (Persona, Día, Desde, Hasta, Puesto)')
    encabezados(ws, CT_CAB, [
        ('A', 'Persona', None), ('B', 'Día', None), ('C', 'Desde (h)', None),
        ('D', 'Hasta (h; la 1:00 es 25)', None), ('E', 'Puesto (L línea, B barra)', None),
        ('F', 'Horas', None), ('G', 'Horas de 22:00 a 6:00 (ET)', None),
        ('H', 'Horas en la franja 1 del plus', None), ('I', 'Horas en la franja 2 del plus', None),
        ('J', 'Horas en la ventana nocturna del convenio', None)], alto=48, congelar=False)
    et_i, et_f, dia = PB('et_ini'), PB('et_fin'), PB('dia')
    p1i, p1f, p2i, p2f = PB('p1_ini'), PB('p1_fin'), PB('p2_ini'), PB('p2_fin')
    for i, (pid, d, a, b, puesto) in enumerate(TURNOS):
        fila = CT_INI + i
        entrada(ws, 'A%d' % fila, pid, etiqueta='persona')
        entrada(ws, 'B%d' % fila, d, etiqueta='día')
        for letra, v in (('C', a), ('D', b)):
            cel = entrada(ws, '%s%d' % (letra, fila), v, fmt=DEC1, etiqueta='hora')
            motor.dv_numerica(ws, [cel.coordinate], minimo=0, maximo=48,
                              mensaje='Hora en decimales: 4,5 son las 4:30 y 25 la 1:00 del día '
                                      'siguiente.')
        entrada(ws, 'E%d' % fila, puesto, etiqueta='puesto')
        A, B = 'C%d' % fila, 'D%d' % fila
        formula(ws, 'F%d' % fila, '=MAX(0,%s-%s)' % (B, A), fmt=DEC1)
        formula(ws, 'G%d' % fila, '=%s+%s' % (_ol(A, B, et_i, et_f + '+' + dia),
                                              _ol(A, B, '0', et_f)), fmt=DEC1)
        formula(ws, 'H%d' % fila, '=%s' % _ol(A, B, p1i, p1f), fmt=DEC1)
        formula(ws, 'I%d' % fila, '=%s+%s' % (_ol(A, B, p2i, p2f),
                                              _ol(A, B, p2i + '+' + dia, p2f + '+' + dia)), fmt=DEC1)
        formula(ws, 'J%d' % fila, '=%s+%s' % (_ol(A, B, p1i, p2f + '+' + dia),
                                              _ol(A, B, p2i, p2f)), fmt=DEC1)
    nota_fuente(ws, 'A%d' % CT_INI, '', extra='Supuestos declarados de El Molinete. '
                'Los tres casos tipo: el churrero/a 1 de 6:00 a 14:00 '
                'entre semana, el churrero/a 2 de 4:30 a 12:30 el fin de semana y quien cierra '
                'el viernes y el sábado de 17:00 a 1:00.')
    nota_legal(ws, 'G%d' % CT_CAB, 'CUN-29')
    nota_legal(ws, 'H%d' % CT_CAB, 'CUN-40')
    nota_legal(ws, 'F%d' % CT_CAB, 'CHN-68')

    # --- rejillas ---------------------------------------------------------------------
    seccion(ws, 'A%d' % CG_SEC, 'Quién Está en Cada Media Hora (y Dónde Falta Gente)')
    grupos = ((G_PRES, 'Personas presentes'), (G_LIN, 'De ellas en la línea'),
              (G_MIN, 'Mínimo exigido (personas)'), (G_HUE, 'Hueco (1 = falta gente)'))
    cab = [('A', 'Desde (h)', None), ('B', 'Hora', None)]
    for g, titulo in grupos:
        ws.merge_cells('%s%d:%s%d' % (g, CG_GRUPO, col(g, 6), CG_GRUPO))
        motor.val(ws, '%s%d' % (g, CG_GRUPO), titulo, bold=True)
        ws['%s%d' % (g, CG_GRUPO)].alignment = Alignment(horizontal='center')
        for j, d in enumerate(DIAS):
            cab.append((col(g, j), d, 6.5))
    encabezados(ws, CG_CAB, cab, alto=22, congelar=False)
    t_per = '$A$%d:$A$%d' % (CT_INI, CT_FIN)
    t_dia = '$B$%d:$B$%d' % (CT_INI, CT_FIN)
    t_ini = '$C$%d:$C$%d' % (CT_INI, CT_FIN)
    t_fin = '$D$%d:$D$%d' % (CT_INI, CT_FIN)
    t_pue = '$E$%d:$E$%d' % (CT_INI, CT_FIN)
    del t_per
    m_ini = '$B$%d:$B$%d' % (CM_INI, CM_FIN)
    m_fin = '$C$%d:$C$%d' % (CM_INI, CM_FIN)
    m_tot = '$D$%d:$D$%d' % (CM_INI, CM_FIN)
    m_lin = '$E$%d:$E$%d' % (CM_INI, CM_FIN)
    for s in range(N_SLOTS):
        fila = CG_INI + s
        if s == 0:
            formula(ws, 'A%d' % fila, '=%s' % PB('slot0'), fmt=DEC1)
        else:
            formula(ws, 'A%d' % fila, '=A%d+%s' % (fila - 1, PB('tramo')), fmt=DEC1)
        formula(ws, 'B%d' % fila, '=IFERROR(A%d/%s,"")' % (fila, dia), fmt=HORA)
        slot = '$A%d' % fila
        for j, d in enumerate(DIAS):
            hdr = '%s$%d' % (col(G_PRES, j), CG_CAB)
            flag = '$%s$%d:$%s$%d' % (col('F', j), CM_INI, col('F', j), CM_FIN)
            pres = 'SUMPRODUCT(--(%s=%s),--(%s<=%s),--(%s>%s))' % (t_dia, hdr, t_ini, slot, t_fin, slot)
            lin = ('SUMPRODUCT(--(%s=%s),--(%s<=%s),--(%s>%s),--(%s="L"))'
                   % (t_dia, hdr, t_ini, slot, t_fin, slot, t_pue))
            req_t = 'SUMPRODUCT(%s,--(%s<=%s),--(%s>%s),%s)' % (flag, m_ini, slot, m_fin, slot, m_tot)
            req_l = 'SUMPRODUCT(%s,--(%s<=%s),--(%s>%s),%s)' % (flag, m_ini, slot, m_fin, slot, m_lin)
            cp, cl, cm, ch = (col(G_PRES, j), col(G_LIN, j), col(G_MIN, j), col(G_HUE, j))
            formula(ws, '%s%d' % (cp, fila), '=' + pres, fmt='0')
            formula(ws, '%s%d' % (cl, fila), '=' + lin, fmt='0')
            formula(ws, '%s%d' % (cm, fila), '=' + req_t, fmt='0')
            formula(ws, '%s%d' % (ch, fila), '=IF(OR(%s%d<%s,%s%d<%s%d),1,0)'
                    % (cl, fila, req_l, cp, fila, cm, fila), fmt='0')
    motor.val(ws, 'A%d' % CG_TOT, 'Huecos por día (medias horas)', bold=True)
    for j in range(len(DIAS)):
        ch = col(G_HUE, j)
        formula(ws, '%s%d' % (ch, CG_TOT), '=SUM(%s%d:%s%d)' % (ch, CG_INI, ch, CG_FIN), fmt='0',
                bold=True)
    motor.regla_expresion(ws, '%s%d:%s%d' % (G_HUE, CG_INI, col(G_HUE, 6), CG_TOT),
                          '=AND(ISNUMBER(%s%d),%s%d>0)' % (G_HUE, CG_INI, G_HUE, CG_INI))

    # --- listas -------------------------------------------------------------------------
    fila_l = CG_TOT + 3
    refs, fin = CC.bloque_listas(ws, fila_l, (('Día', DIAS), ('Puesto', ('L', 'B'))))
    CC.dv_rango(ws, ['B%d' % r for r in range(CT_INI, CT_FIN + 1)], refs['Día'], 'Día',
                'Elige L, M, X, J, V, S o D.')
    CC.dv_rango(ws, ['E%d' % r for r in range(CT_INI, CT_FIN + 1)], refs['Puesto'], 'Puesto',
                'L = línea de fritura; B = barra, sala y despacho.')
    CC.dv_rango(ws, ['A%d' % r for r in range(CT_INI, CT_FIN + 1)],
                '=' + RNG(H_COS, 'A%d' % K_INI, 'A%d' % K_FIN), 'Persona',
                'Elige una persona de «Coste de Plantilla».')
    version_al_pie(ws, fin + 1)
    setup(ws, apaisado=True)
    ws.freeze_panes = None
    return ws


# ==========================================================================
# Hoja «Nocturnidad ET y Convenio»
# ==========================================================================
VER_SI = 'sí'
VER_NO = 'no'
VER_REV = 'revísalo con tu asesor'
VER_NA = 'No aplica (autónomo)'
QUE_TOCA = {
    VER_SI: ('Trabajador nocturno: no puede hacer horas extra y su jornada no pasa de ocho '
             'horas diarias de media'),
    VER_NO: ('No es trabajador nocturno: el Estatuto no le pone esos límites; el plus del '
             'convenio se paga igual'),
    VER_REV: ('Revísalo con tu asesor: no llega a tres horas todos los días, pero hay días en '
              'que sí; el «normalmente» del artículo no está interpretado'),
}
QUE_TOCA_TITULAR = ('Titular autónomo: ni las dos figuras ni el plus le alcanzan; sus horas, en '
                    '«Horas del Titular»')
AVISO_JORNADA = ('Hay días con la jornada entera nocturna según el convenio: este libro paga '
                 'el plus hora a hora; revísalo con tu asesor')


def hoja_nocturnidad(wb):
    ws = wb.create_sheet(H_NOC)
    cabecera_hoja(ws, H_NOC, 'Dos columnas por persona: la del Estatuto (¿es trabajador '
                             'nocturno?) y la del convenio (¿cuánto plus cobra?). Son dos '
                             'figuras distintas y aquí no se mezclan.')
    for letra, ancho in (('A', 9), ('B', 26), ('C', 11), ('D', 10), ('E', 11), ('F', 11),
                         ('G', 22), ('H', 11), ('I', 11), ('J', 12), ('K', 12), ('L', 12),
                         ('M', 40), ('N', 34)):
        ws.column_dimensions[letra].width = ancho
    ws.merge_cells('C%d:G%d' % (N_GRUPO, N_GRUPO))
    motor.val(ws, 'C%d' % N_GRUPO, 'FIGURA 1 · Trabajador nocturno del Estatuto (art. 36.1): '
                                   'de 22:00 a 6:00', bold=True, wrap=True)
    nota_legal(ws, 'C%d' % N_GRUPO, 'CUN-29')
    ws.merge_cells('H%d:L%d' % (N_GRUPO, N_GRUPO))
    motor.val(ws, 'H%d' % N_GRUPO, 'FIGURA 2 · Plus de nocturnidad del convenio de ejemplo',
              bold=True, wrap=True)
    nota_legal(ws, 'H%d' % N_GRUPO, 'CUN-40')
    ws.row_dimensions[N_GRUPO].height = 30
    encabezados(ws, N_CAB, [
        ('A', 'Persona', None), ('B', 'Puesto', None),
        ('C', 'Horas de 22:00 a 6:00 a la semana', None), ('D', 'Días que trabaja', None),
        ('E', 'Días con 3 h o más en ese periodo', None), ('F', 'Parte de su jornada en ese periodo', None),
        ('G', '¿Trabajador nocturno? (ET)', None),
        ('H', 'Horas en la franja 1 a la semana', None), ('I', 'Horas en la franja 2 a la semana', None),
        ('J', 'Días con jornada entera nocturna', None), ('K', 'Plus a la semana', None),
        ('L', 'Plus al año', None), ('M', 'Qué le toca por la figura 1', None),
        ('N', 'Aviso de la figura 2', None)], alto=54, congelar=False)

    t_per = RNG(H_CUA, 'A%d' % CT_INI, 'A%d' % CT_FIN)
    t_dia = RNG(H_CUA, 'B%d' % CT_INI, 'B%d' % CT_FIN)
    t_col = dict((k, RNG(H_CUA, '%s%d' % (k, CT_INI), '%s%d' % (k, CT_FIN))) for k in 'FGHIJ')
    # --- matrices por día ---------------------------------------------------------------
    matrices = ((N_M1_CAB, N_M1, 'F', 'Horas trabajadas por día'),
                (N_M2_CAB, N_M2, 'G', 'Horas de 22:00 a 6:00 por día (Estatuto)'),
                (N_M3_CAB, N_M3, 'J', 'Horas en la ventana nocturna del convenio por día'))
    seccion(ws, 'A%d' % (N_M1_CAB - 1), 'Horas por Día (Lo que Usan las Columnas de Arriba)')
    for cab, ini, k, titulo in matrices:
        encabezados(ws, cab, [('A', 'Persona', None), ('B', titulo, None)]
                    + [(col('C', j), d, None) for j, d in enumerate(DIAS)], alto=30, congelar=False)
        for i in range(N_P):
            fila = ini + i
            formula(ws, 'A%d' % fila, '=%s' % ABS_(H_COS, 'A%d' % (K_INI + i)))
            for j in range(len(DIAS)):
                c = col('C', j)
                formula(ws, '%s%d' % (c, fila), '=SUMPRODUCT(--(%s=$A%d),--(%s=%s$%d),%s)'
                        % (t_per, fila, t_dia, c, cab, t_col[k]), fmt=DEC1)
    # --- la tabla de las dos figuras --------------------------------------------------------
    et_min, tercio, umbral = PB('et_min'), PB('et_tercio'), PB('umbral')
    for i in range(N_P):
        fila = N_INI + i
        kf = K_INI + i
        m1 = 'C%d:I%d' % (N_M1 + i, N_M1 + i)
        m2 = 'C%d:I%d' % (N_M2 + i, N_M2 + i)
        m3 = 'C%d:I%d' % (N_M3 + i, N_M3 + i)
        formula(ws, 'A%d' % fila, '=%s' % ABS_(H_COS, 'A%d' % kf), bold=True)
        formula(ws, 'B%d' % fila, '=%s' % ABS_(H_COS, 'B%d' % kf))
        formula(ws, 'C%d' % fila, '=SUM(%s)' % m2, fmt=DEC1)
        formula(ws, 'D%d' % fila, '=COUNTIF(%s,">0")' % m1, fmt=ENT)
        formula(ws, 'E%d' % fila, '=COUNTIF(%s,">="&%s)' % (m2, et_min), fmt=ENT)
        formula(ws, 'F%d' % fila, '=IFERROR(C%d/%s,"")' % (fila, ABS_(H_COS, 'G%d' % kf)), fmt=PCT)
        # R-20: al titular, que es autónomo, no se le aplica el Estatuto: «No aplica (autónomo)»
        formula(ws, 'G%d' % fila,
                '=IF({tit}="{sitit}","{na}",IF(AND(ISNUMBER(D{f}),ISNUMBER(E{f}),ISNUMBER(F{f})),'
                'IF(AND(E{f}=0,F{f}<{t}),"{no}",IF(OR(AND(D{f}>0,E{f}=D{f}),F{f}>={t}),"{si}",'
                '"{rev}")),""))'
                .format(f=fila, t=tercio, no=VER_NO, si=VER_SI, rev=VER_REV, na=VER_NA,
                        tit=ABS_(H_COS, 'E%d' % kf), sitit=SI), bold=True)
        semaforo(ws, 'G%d' % fila, (VER_NO,), (VER_SI,), (VER_REV,))
        formula(ws, 'H%d' % fila, '=SUMPRODUCT(--(%s=$A%d),%s)' % (t_per, fila, t_col['H']), fmt=DEC1)
        formula(ws, 'I%d' % fila, '=SUMPRODUCT(--(%s=$A%d),%s)' % (t_per, fila, t_col['I']), fmt=DEC1)
        formula(ws, 'J%d' % fila, '=COUNTIF(%s,">="&%s)' % (m3, umbral), fmt=ENT)
        formula(ws, 'K%d' % fila, '=IF(%s="%s",0,(H%d*%s+I%d*%s)*%s)'
                % (ABS_(H_COS, 'E%d' % kf), SI, fila, PB('p1_pct'), fila, PB('p2_pct'),
                   ABS_(H_COS, 'L%d' % kf)), fmt=EUR)
        formula(ws, 'L%d' % fila, '=K%d*%s' % (fila, PB('sem_tr')), fmt=EUR, bold=True, destacar=True)
        formula(ws, 'M%d' % fila, '=IF({tit}="{si}","{t}",IF(G{f}="{si}","{a}",IF(G{f}="{no}","{b}",'
                                  'IF(G{f}="{r}","{c}",""))))'
                .format(f=fila, si=VER_SI, no=VER_NO, r=VER_REV, a=QUE_TOCA[VER_SI],
                        b=QUE_TOCA[VER_NO], c=QUE_TOCA[VER_REV], tit=ABS_(H_COS, 'E%d' % kf),
                        t=QUE_TOCA_TITULAR))
        ws['M%d' % fila].alignment = Alignment(wrap_text=True, vertical='top')
        formula(ws, 'N%d' % fila, '=IF(AND(ISNUMBER(J%d),J%d>0),"%s","")' % (fila, fila, AVISO_JORNADA))
        ws['N%d' % fila].alignment = Alignment(wrap_text=True, vertical='top')
        ws.row_dimensions[fila].height = 58
    notas = [
        'Figura 1. Las horas del periodo nocturno del Estatuto se cuentan día a día en la tabla de '
        'abajo. «Sí» si todos los días que trabaja hace tres horas o más en ese periodo, o si ese '
        'periodo es un tercio de su jornada o más; «no» si ningún día llega a tres y no llega al '
        'tercio; «revísalo con tu asesor» en medio.',
        'Figura 2. El plus se paga por hora sobre el salario base por hora del puesto («Coste de '
        'Plantilla», columna L), con el recargo de cada franja. El titular no lo cobra: es '
        'autónomo. Si tu convenio paga la jornada entera cuando pasa de su umbral, el aviso de la '
        'columna N te lo dice.',
        'Ejemplo de El Molinete: el turno de las 4:30 del fin de semana tiene hora y media en el '
        'periodo del Estatuto (no es trabajador nocturno) y tres horas y media con plus. Quien '
        'cierra el viernes y el sábado hasta la 1:00 tiene tres horas dos días por semana: '
        '«revísalo con tu asesor».',
    ]
    for i, txt in enumerate(notas):
        fila = N_NOTAS + i
        ws.merge_cells('A%d:N%d' % (fila, fila))
        motor.val(ws, 'A%d' % fila, txt, wrap=True)
        ws['A%d' % fila].font = Font(size=9, italic=True)
        ws.row_dimensions[fila].height = 32
    version_al_pie(ws, N_M3 + N_P + 2)
    setup(ws, apaisado=True)
    return ws


# ==========================================================================
# Hoja «Coste de Plantilla»
# ==========================================================================
SMI_BAJO = 'Por debajo del SMI 2026: cuenta el SMI'
SMI_OK = 'Por encima del SMI 2026'
JOR_PASA = 'Pasa de su jornada'
JOR_OK = 'Dentro de su jornada'
JOR_TIT = 'Titular: sus horas de más, en «Horas del Titular»'


def hoja_coste(wb):
    ws = wb.create_sheet(H_COS)
    cabecera_hoja(ws, H_COS, 'Coste anual por puesto = MAX(salario; SMI anual a su jornada) más '
                             'la Seguridad Social de empresa. Convenio de ejemplo: Madrid, clase '
                             'C, tablas de 2025 (vencido, en negociación).')
    seccion(ws, 'K4', 'Lo que Copian Otros Libros (No Muevas Estas Celdas)')
    motor.val(ws, 'K5', 'Coste anual de plantilla (puestos, refuerzo y sustituciones)', wrap=True)
    ws['K5'].font = Font(bold=True, size=9)
    formula(ws, 'L5', '=B%d' % K_TOTAL, fmt=EUR, bold=True, destacar=True)
    c = CRUCES_OUT[0]
    gris(ws, 'M5', 'Lo copia: libro %d (%s, hoja %s, %s)' % (c['receptor'], c['fichero_receptor'],
                                                             c['hoja_receptor'], c['x']))
    ws.row_dimensions[5].height = 30

    cols = [('A', 'Persona', 9), ('B', 'Puesto', 28), ('C', 'Jornada (1,0 = completa)', 11),
            ('D', 'Nivel del convenio', 9), ('E', '¿Es el titular?', 9),
            ('F', '¿Cotiza Seguridad Social de empresa?', 11),
            ('G', 'Horas a la semana en el cuadrante', 11), ('H', 'Horas de su jornada', 10),
            ('I', '¿Cabe en su jornada?', 20), ('J', 'Salario base mensual', 12),
            ('K', 'Bruto anual a su jornada', 46), ('L', 'Salario base por hora', 14),
            ('M', 'Plus de nocturnidad al año', 46), ('N', 'Salario anual', 13),
            ('O', 'Suelo: SMI anual a su jornada', 13), ('P', 'Semáforo del SMI', 22),
            ('Q', 'Coste anual del puesto', 14)]
    encabezados(ws, K_CAB, cols, alto=54, congelar=False)
    niv_keys = RNG(H_PAR, 'A%d' % min(NIV_ROWS.values()), 'A%d' % max(NIV_ROWS.values()))
    niv_vals = RNG(H_PAR, 'B%d' % min(NIV_ROWS.values()), 'B%d' % max(NIV_ROWS.values()))
    t_per = RNG(H_CUA, 'A%d' % CT_INI, 'A%d' % CT_FIN)
    t_h = RNG(H_CUA, 'F%d' % CT_INI, 'F%d' % CT_FIN)
    for i, p in enumerate(PLANTILLA):
        fila = K_INI + i
        motor.val(ws, 'A%d' % fila, p['id'], bold=True)
        entrada(ws, 'B%d' % fila, p['puesto'], etiqueta='puesto')
        nota_fuente(ws, 'B%d' % fila, '', extra=p['nota'])
        cel = entrada(ws, 'C%d' % fila, p['jornada'], fmt=DEC, etiqueta='jornada')
        motor.dv_numerica(ws, [cel.coordinate], minimo=0, maximo=1,
                          mensaje='Jornada en tanto por uno: 1 = completa, 0,5 = media.')
        entrada(ws, 'D%d' % fila, p['nivel'], etiqueta='nivel')
        entrada(ws, 'E%d' % fila, SI if p['es_titular'] else NO, etiqueta='titular')
        entrada(ws, 'F%d' % fila, SI if p['cotiza_ss_empresa'] else NO, etiqueta='cotiza')
        formula(ws, 'G%d' % fila, '=SUMPRODUCT(--(%s=A%d),%s)' % (t_per, fila, t_h), fmt=DEC1)
        formula(ws, 'H%d' % fila, '=C%d*%s' % (fila, PB('h_sem')), fmt=DEC1)
        formula(ws, 'I%d' % fila, '=IF(E{f}="{si}","{tit}",IF(AND(ISNUMBER(G{f}),ISNUMBER(H{f})),'
                                  'IF(G{f}>H{f},"{pasa}","{ok}"),""))'
                .format(f=fila, si=SI, tit=JOR_TIT, pasa=JOR_PASA, ok=JOR_OK))
        semaforo(ws, 'I%d' % fila, (JOR_OK,), (JOR_PASA,), (JOR_TIT,))
        formula(ws, 'J%d' % fila, '=SUMIF(%s,D%d,%s)' % (niv_keys, fila, niv_vals), fmt=EUR)
        formula(ws, 'K%d' % fila, '=(J%d*%s+%s*%s)*C%d' % (fila, PB('pagas'), PB('plus_conv'),
                                                            PB('meses_plus'), fila), fmt=EUR)
        formula(ws, 'L%d' % fila, '=IFERROR(J%d*%s/%s,"")' % (fila, PB('pagas'), PB('jor_anual')),
                fmt=EUR)
        formula(ws, 'M%d' % fila, '=%s' % ABS_(H_NOC, 'L%d' % (N_INI + i)), fmt=EUR)
        formula(ws, 'N%d' % fila, '=K%d+M%d' % (fila, fila), fmt=EUR)
        formula(ws, 'O%d' % fila, '=%s*C%d' % (PB('smi_anual'), fila), fmt=EUR)
        formula(ws, 'P%d' % fila, '=IF(AND(ISNUMBER(N{f}),ISNUMBER(O{f})),IF(N{f}<O{f},"{b}","{o}"),"")'
                .format(f=fila, b=SMI_BAJO, o=SMI_OK))
        semaforo(ws, 'P%d' % fila, (SMI_OK,), (SMI_BAJO,))
        formula(ws, 'Q%d' % fila, '=IF(F{f}="{si}",MAX(N{f},O{f})*(1+{ss}),MAX(N{f},O{f}))'
                .format(f=fila, si=SI, ss=PB('ss')), fmt=EUR, bold=True, destacar=True)
        ws.row_dimensions[fila].height = 30
    nota_legal(ws, 'O%d' % K_CAB, 'CHN-67')
    nota_legal(ws, 'J%d' % K_CAB, 'CUN-39')
    nota_legal(ws, 'M%d' % K_CAB, 'CUN-40')
    motor.val(ws, 'A%d' % K_TOT, 'Total', bold=True)
    formula(ws, 'C%d' % K_TOT, '=SUM(C%d:C%d)' % (K_INI, K_FIN), fmt=DEC, bold=True)
    formula(ws, 'G%d' % K_TOT, '=SUM(G%d:G%d)' % (K_INI, K_FIN), fmt=DEC1, bold=True)
    formula(ws, 'M%d' % K_TOT, '=SUM(M%d:M%d)' % (K_INI, K_FIN), fmt=EUR, bold=True)
    formula(ws, 'Q%d' % K_TOT, '=SUM(Q%d:Q%d)' % (K_INI, K_FIN), fmt=EUR, bold=True, destacar=True)

    # --- del coste de los puestos al coste de plantilla ------------------------------------
    seccion(ws, 'A%d' % K_SEC, 'Del Coste de los Puestos al Coste Anual de Plantilla')
    filas = (
        (K_PUESTOS, 'Coste de los puestos al año', '=Q%d' % K_TOT, EUR),
        (K_HORAS, 'Horas contratadas al año (jornadas por la jornada anual)',
         '=C%d*%s' % (K_TOT, PB('jor_anual')), DEC1),
        (K_CHORA, 'Coste hora de la plantilla (el que tiene que usar el libro 3)',
         '=IFERROR(B%d/B%d,"")' % (K_PUESTOS, K_HORAS), EUR),
        (K_HEMP, 'Horas a la semana de los empleados (sin el titular)',
         '=SUMIF(E%d:E%d,"%s",G%d:G%d)' % (K_INI, K_FIN, NO, K_INI, K_FIN), DEC1),
        (K_SEMS, 'Semanas que el local abre y un contrato no trabaja',
         '=MAX(0,%s-%s)' % (PB('sem_ab'), PB('sem_tr')), DEC),
        (K_HSUS, 'Horas de sustitución al año (vacaciones y festivos)',
         '=IF(%s="%s",B%d*B%d,0)' % (PB('cubrir'), SI, K_HEMP, K_SEMS), DEC1),
        (K_HREF, 'Horas de refuerzo de Navidad y Reyes al año (copiadas del libro 5)',
         '=%s' % XB('horas_refuerzo_pico'), DEC1),
        (K_CEXT, 'Coste de sustituciones y refuerzo al coste hora',
         '=(B%d+B%d)*B%d' % (K_HSUS, K_HREF, K_CHORA), EUR),
        (K_TOTAL, 'COSTE ANUAL DE PLANTILLA (lo copia el libro 6)',
         '=B%d+B%d' % (K_PUESTOS, K_CEXT), EUR),
    )
    for fila, etq, form, fmt in filas:
        motor.val(ws, 'A%d' % fila, etq, wrap=True)
        formula(ws, 'B%d' % fila, form, fmt=fmt, bold=True, destacar=fila in (K_CHORA, K_TOTAL))
        ws.row_dimensions[fila].height = 28
    gris(ws, 'C%d' % K_CHORA, 'Cuadra con el coste hora que usaste en el escandallo del libro 3 '
                              '(fila de CUADRE de «Parámetros»).')
    gris(ws, 'C%d' % K_HSUS, 'Cada contrato trabaja sus horas anuales, no las semanas que abre '
                             'el local: las semanas que faltan las cubre un sustituto.')

    # --- salarios de mercado -------------------------------------------------------------
    seccion(ws, 'A%d' % K_MER_SEC, 'Salarios de Mercado (Sólo Como Referencia)')
    gris(ws, 'A%d' % (K_MER_SEC + 1), D.NOTA_SALARIOS_MERCADO, alto=28)
    ws.merge_cells('A%d:I%d' % (K_MER_SEC + 1, K_MER_SEC + 1))
    encabezados(ws, K_MER_CAB, [('A', 'Fuente', None), ('B', 'Oferta publicada', None),
                                ('C', 'Mínimo EUR/mes', None), ('D', 'Máximo EUR/mes', None),
                                ('E', 'Mínimo al año con tus pagas de lectura', None),
                                ('F', 'Frente al SMI anual', None)], alto=42, congelar=False)
    for i, (puesto, mn, mx, fuente, nota_m) in enumerate(D.SALARIOS_MERCADO):
        fila = K_MER_INI + i
        motor.val(ws, 'A%d' % fila, fuente)
        motor.val(ws, 'B%d' % fila, puesto, wrap=True)
        motor.val(ws, 'C%d' % fila, mn, fmt=EUR)
        motor.val(ws, 'D%d' % fila, mx, fmt=EUR)
        nota_fuente(ws, 'B%d' % fila, fuente, extra=nota_m)
        formula(ws, 'E%d' % fila, '=C%d*%s' % (fila, PB('pagas_oferta')), fmt=EUR)
        formula(ws, 'F%d' % fila, '=IF(ISNUMBER(E{f}),IF(E{f}<{s},"{b}","{o}"),"")'
                .format(f=fila, s=PB('smi_anual'), b=SMI_BAJO, o=SMI_OK))
        semaforo(ws, 'F%d' % fila, (SMI_OK,), (SMI_BAJO,))
        ws.row_dimensions[fila].height = 30

    refs, fin = CC.bloque_listas(ws, K_MER_FIN + 3, (('Sí o no', SI_NO),))
    CC.dv_rango(ws, ['E%d' % r for r in range(K_INI, K_FIN + 1)]
                + ['F%d' % r for r in range(K_INI, K_FIN + 1)], refs['Sí o no'], 'Sí o no',
                'Elige sí o no.')
    CC.dv_rango(ws, ['D%d' % r for r in range(K_INI, K_FIN + 1)], '=' + niv_keys, 'Nivel',
                'Elige un nivel de la tabla de «Parámetros».')
    version_al_pie(ws, fin + 1)
    setup(ws, apaisado=True)
    ws.freeze_panes = None
    return ws


# ==========================================================================
# Hoja «Identifica tu Convenio»
# ==========================================================================
PASOS_CONVENIO = [
    ('1', 'Tu provincia', 'Escríbela en «Parámetros». El convenio de hostelería es provincial o '
     'autonómico.', None, 'provincia', None),
    ('2', 'Busca el convenio de hostelería de tu provincia', 'En REGCON, por denominación y '
     'ámbito; anota su código.', 'CHN-93', 'codigo', None),
    ('3', 'Comprueba en qué clase o grupo entra tu local', 'Lee su ámbito y su clasificación de '
     'establecimientos. En Madrid, la clase C agrupa cafeterías, chocolaterías y bares, y el '
     'convenio nombra freidurías, quioscos y chocolaterías.', 'CUN-39', 'clase', None),
    ('4', 'Comprueba si está vigente', 'Si venció y está en negociación, se aplican sus tablas '
     'mientras no se publique el nuevo. Mira el boletín antes de cada versión de tus números.',
     'CUN-V01', 'tablas', None),
    ('5', 'Encaja cada puesto en un nivel por sus funciones', 'Es una asimilación de funciones: '
     'compruébala con tu asesor. El acuerdo laboral estatal de hostelería nombra freidurías, '
     'quioscos y chocolaterías.', 'CUN-31', None,
     'Línea de fritura al nivel III, sala al IV y refuerzo al V (asimilación propuesta)'),
    ('6', 'Comprueba hasta cuándo rige el acuerdo estatal', 'El VI Acuerdo Laboral estatal de '
     'hostelería sigue vigente hasta el 31-12-2030.', 'CUN-32', None, 'Vigente hasta el 31-12-2030'),
    ('7', 'Si tienes obrador, mira si otro convenio lo alcanza', 'Ejemplo del método: el '
     'convenio de Toledo nombra los obradores de masas fritas. Alcanza a un obrador; si lo tuyo '
     'es un despacho con sala, te lo confirma un laboralista.', 'CUN-V05', None,
     'No aplica: despacho con sala'),
    ('8', 'Anota el plus de nocturnidad y sus franjas', 'Cópialo a «Parámetros», figura 2.',
     'CUN-40', None, 'De 22:00 a 0:00 +1 %; de 0:00 a 8:00 +25 %'),
    ('9', 'Organiza el registro de jornada', 'Diario, con hora de inicio y de fin, guardado '
     'cuatro años.', 'CHN-68', None, 'Registro diario en papel firmado'),
]


def hoja_identifica(wb):
    ws = wb.create_sheet(H_IDE)
    cabecera_hoja(ws, H_IDE, 'El método para encontrar TU convenio, con el de Madrid como '
                             'ejemplo. Nada de esta hoja sustituye a tu asesor laboral.')
    encabezados(ws, 5, [('A', 'Paso', 6), ('B', 'Qué haces', 34), ('C', 'Cómo', 60),
                        ('D', 'Tu respuesta', 46)], alto=24, congelar=False)
    fila = 5
    for paso, que, como, legal, clave_par, defecto in PASOS_CONVENIO:
        fila += 1
        motor.val(ws, 'A%d' % fila, paso, bold=True)
        motor.val(ws, 'B%d' % fila, que, wrap=True, bold=True)
        motor.val(ws, 'C%d' % fila, como, wrap=True)
        ws['C%d' % fila].font = Font(size=9)
        if clave_par:
            formula(ws, 'D%d' % fila, '=%s' % PB(clave_par))
            ws['D%d' % fila].alignment = Alignment(wrap_text=True, vertical='top')
            gris(ws, 'E%d' % fila, 'Se escribe en «Parámetros»!B%d.' % PR[clave_par])
        else:
            entrada(ws, 'D%d' % fila, defecto, etiqueta='respuesta del paso ' + paso)
            ws['D%d' % fila].alignment = Alignment(wrap_text=True, vertical='top')
        if legal:
            nota_legal(ws, 'C%d' % fila, legal)
        ws.row_dimensions[fila].height = 48
    ws.column_dimensions['E'].width = 30
    fila += 2
    seccion(ws, 'A%d' % fila, 'El Ejemplo de Madrid, de un Vistazo')
    M = D.CONVENIO_METODO
    ejemplo = [
        ('Convenio', '=%s' % PB('denominacion')),
        ('Código REGCON', '=%s' % PB('codigo')),
        ('Clase', '=%s' % PB('clase')),
        ('Tablas', '=%s' % PB('tablas')),
        ('Pagas (supuesto)', '=%s' % PB('pagas')),
        ('Método con otro convenio (ejemplo)', 'Toledo, %s: %s (%s)' % (
            D.dato(M, 'toledo_codigo'), D.dato(M, 'toledo_ambito'), D.dato(M, 'toledo_vigencia'))),
    ]
    for etq, v in ejemplo:
        fila += 1
        motor.val(ws, 'B%d' % fila, etq, bold=True)
        if isinstance(v, str) and v.startswith('='):
            formula(ws, 'C%d' % fila, v)
        else:
            motor.val(ws, 'C%d' % fila, v, wrap=True)
        ws['C%d' % fila].alignment = Alignment(wrap_text=True, vertical='top')
        ws.row_dimensions[fila].height = 30
        if etq.startswith('Método'):
            nota_legal(ws, 'C%d' % fila, 'CUN-V05')
    fila += 2
    ws.merge_cells('A%d:D%d' % (fila, fila))
    motor.val(ws, 'A%d' % fila, D.dato(M, 'como_buscar') + '.', wrap=True)
    ws['A%d' % fila].font = Font(size=9, italic=True)
    nota_legal(ws, 'A%d' % fila, 'CHN-93')
    version_al_pie(ws, fila + 2)
    setup(ws, apaisado=True)
    return ws


# ==========================================================================
# Hoja «Horas del Titular»
# ==========================================================================
TIT_7 = ('Trabajas los siete días: un plan que sólo cuadra con el titular sin descanso no '
         'aguanta un año; busca un relevo para un día')
TIT_EXTRA = 'Haces más de una jornada completa: esas horas no se pagan, pero existen'
TIT_OK = 'Tus horas caben en una jornada completa con descanso semanal'


def hoja_titular(wb):
    ws = wb.create_sheet(H_TIT)
    cabecera_hoja(ws, H_TIT, 'Las horas que haces tú. No se pagan aparte (tu retribución ya va '
                             'en el coste de plantilla), pero existen, y son la decisión 10: '
                             'cuánta madrugada asumes tú.')
    for letra, ancho in (('A', 60), ('B', 18), ('C', 60)):
        ws.column_dimensions[letra].width = ancho
    encabezados(ws, 5, [('A', 'Concepto', None), ('B', 'Valor', None), ('C', 'Nota', None)],
                alto=24, congelar=False)
    e = RNG(H_COS, 'E%d' % K_INI, 'E%d' % K_FIN)

    def tit(rango_hoja, c1, c2):
        return 'SUMIF(%s,"%s",%s)' % (e, SI, RNG(rango_hoja, c1, c2))

    filas = (
        (T_HSEM, 'Horas a la semana en el cuadrante',
         '=' + tit(H_COS, 'G%d' % K_INI, 'G%d' % K_FIN), DEC1, ''),
        (T_JOR, 'Horas de una jornada completa', '=%s' % PB('h_sem'), DEC1, ''),
        (T_EXTRA, 'Horas de más a la semana', '=MAX(0,B%d-B%d)' % (T_HSEM, T_JOR), DEC1, ''),
        (T_DIAS, 'Días que trabajas a la semana',
         '=' + tit(H_NOC, 'D%d' % N_INI, 'D%d' % N_FIN), ENT, ''),
        (T_MADRU, 'Horas de madrugada a la semana (franja 2 del plus: de 0:00 a 8:00)',
         '=' + tit(H_NOC, 'I%d' % N_INI, 'I%d' % N_FIN), DEC1,
         'En El Molinete, las aperturas de las 6:00 del sábado y el domingo y la primera hora '
         'del lunes.'),
        (T_SEMS, 'Semanas que el local abre al año', '=%s' % PB('sem_ab'), DEC, ''),
        (T_HANIO, 'Horas al año', '=B%d*B%d' % (T_HSEM, T_SEMS), DEC1, ''),
        (T_EXANIO, 'Horas de más al año', '=B%d*B%d' % (T_EXTRA, T_SEMS), DEC1, ''),
        (T_VALOR, 'Lo que valen tus horas de más al coste hora de la plantilla',
         '=B%d*%s' % (T_EXANIO, ABS_(H_COS, 'B%d' % K_CHORA)), EUR,
         'Lo que costaría pagárselas a otra persona. No está en el coste de plantilla.'),
        (T_VERED, 'Veredicto',
         '=IF(AND(ISNUMBER(B{d}),ISNUMBER(B{x})),IF(B{d}>={s},"{a}",IF(B{x}>0,"{b}","{c}")),"")'
         .format(d=T_DIAS, x=T_EXTRA, s=PB('dsem'), a=TIT_7, b=TIT_EXTRA, c=TIT_OK), None, ''),
    )
    for fila, etq, form, fmt, nota_t in filas:
        motor.val(ws, 'A%d' % fila, etq, wrap=True)
        formula(ws, 'B%d' % fila, form, fmt=fmt, bold=True, destacar=fila in (T_EXANIO, T_VALOR))
        if nota_t:
            gris(ws, 'C%d' % fila, nota_t)
        ws.row_dimensions[fila].height = 28
    ws.merge_cells('B%d:C%d' % (T_VERED, T_VERED))
    ws['B%d' % T_VERED].alignment = Alignment(wrap_text=True, vertical='top')
    ws.row_dimensions[T_VERED].height = 40
    semaforo(ws, 'B%d' % T_VERED, (TIT_OK,), (TIT_7,), (TIT_EXTRA,))
    dec = [d for d in D.DECISIONES if d[0] == 10][0]
    motor.val(ws, 'A%d' % T_DEC, 'Decisión %d: %s' % (dec[0], dec[1]), bold=True)
    motor.val(ws, 'B%d' % T_DEC, 'El Molinete: %s.' % dec[3].lower(), wrap=True)
    ws.merge_cells('B%d:C%d' % (T_DEC, T_DEC))
    ws.row_dimensions[T_DEC].height = 30
    fila = T_DEC + 2
    ws.merge_cells('A%d:C%d' % (fila, fila))
    motor.val(ws, 'A%d' % fila,
              'El titular de El Molinete es autónomo y está en nómina en el plan como un puesto '
              'más (al nivel III de la tabla de ejemplo): un plan que sólo cuadra si el titular '
              'trabaja gratis no sirve. Su cuota de autónomos va en los gastos fijos del libro 6.',
              wrap=True)
    ws['A%d' % fila].font = Font(size=9, italic=True)
    ws.row_dimensions[fila].height = 42
    version_al_pie(ws, fila + 2)
    setup(ws, apaisado=False)
    return ws


# ==========================================================================
# Mapa de celdas
# ==========================================================================
def mapa_celdas():
    m = []
    for c in CRUCES_IN:
        clave = c['calcula']
        m.append(('%s (copiado del libro %d, %s)' % (ETIQUETA_CRUCE[clave][0], c['origen'], c['x']),
                  H_PAR, 'B%d' % CRUCE_ROW[clave], 'entrada'))
    m += [
        ('Coste anual de plantilla (origen de X15)', H_COS, 'L5', 'salida'),
        ('Coste anual de plantilla (total de la hoja)', H_COS, 'B%d' % K_TOTAL, 'salida'),
        ('Coste de los puestos al año', H_COS, 'B%d' % K_PUESTOS, 'salida'),
        ('Coste hora de la plantilla', H_COS, 'B%d' % K_CHORA, 'salida'),
        ('Horas contratadas al año', H_COS, 'B%d' % K_HORAS, 'salida'),
        ('Horas de sustitución al año', H_COS, 'B%d' % K_HSUS, 'salida'),
        ('Coste de sustituciones y refuerzo', H_COS, 'B%d' % K_CEXT, 'salida'),
        ('Plus de nocturnidad de la plantilla al año', H_COS, 'M%d' % K_TOT, 'salida'),
        ('Provincia del convenio', H_PAR, 'B%d' % PR['provincia'], 'entrada'),
        ('Código REGCON del convenio de ejemplo', H_PAR, 'B%d' % PR['codigo'], 'parametro'),
        ('Tablas salariales del ejemplo', H_PAR, 'B%d' % PR['tablas'], 'parametro'),
        ('Pagas al año', H_PAR, 'B%d' % PR['pagas'], 'parametro'),
        ('Plus de convenio al mes', H_PAR, 'B%d' % PR['plus_conv'], 'parametro'),
        ('SMI anual', H_PAR, 'B%d' % PR['smi_anual'], 'parametro'),
        ('Seguridad Social de empresa', H_PAR, 'B%d' % PR['ss'], 'parametro'),
        ('Jornada anual de un contrato completo', H_PAR, 'B%d' % PR['jor_anual'], 'parametro'),
        ('Recargo del plus de 22:00 a 0:00', H_PAR, 'B%d' % PR['p1_pct'], 'parametro'),
        ('Recargo del plus de 0:00 a 8:00', H_PAR, 'B%d' % PR['p2_pct'], 'parametro'),
        ('Umbral de jornada nocturna del convenio', H_PAR, 'B%d' % PR['umbral'], 'parametro'),
        ('Semanas que el local abre al año', H_PAR, 'B%d' % PR['sem_ab'], 'parametro'),
        ('Cuadre del coste hora del escandallo (veredicto)', H_PAR,
         'D%d' % PR['cuadre_coste_hora_mano_obra'], 'salida'),
        ('Contraste del coste hora con tu plantilla (veredicto)', H_PAR,
         'B%d' % PR['contraste_coste_hora_mano_obra'], 'salida'),
        ('Horas de apertura con mínimo de cobertura a la semana', H_CUA, 'B%d' % CR_HAP, 'salida'),
        ('Horas de trabajo que piden los mínimos a la semana', H_CUA, 'B%d' % CR_HEX, 'salida'),
        ('Horas de trabajo que pone el cuadrante a la semana', H_CUA, 'B%d' % CR_HCU, 'salida'),
        ('Huecos de cobertura (medias horas)', H_CUA, 'B%d' % CR_HUE, 'salida'),
        ('Veredicto de cobertura del cuadrante', H_CUA, 'B%d' % CR_VER, 'salida'),
        ('Máximo de personas en la línea a la vez', H_CUA, 'B%d' % CR_MAXL, 'salida'),
        ('Horas del titular a la semana', H_TIT, 'B%d' % T_HSEM, 'salida'),
        ('Horas de más del titular a la semana', H_TIT, 'B%d' % T_EXTRA, 'salida'),
        ('Horas de madrugada del titular a la semana', H_TIT, 'B%d' % T_MADRU, 'salida'),
        ('Horas de más del titular al año', H_TIT, 'B%d' % T_EXANIO, 'salida'),
        ('Valor de las horas de más del titular al año', H_TIT, 'B%d' % T_VALOR, 'salida'),
        ('Veredicto de las horas del titular', H_TIT, 'B%d' % T_VERED, 'salida'),
    ]
    for i, p in enumerate(PLANTILLA):
        quien = p['puesto']
        m.append(('Coste anual del puesto: %s' % quien, H_COS, 'Q%d' % (K_INI + i), 'salida'))
        m.append(('Salario anual (antes del suelo): %s' % quien, H_COS, 'N%d' % (K_INI + i), 'salida'))
        m.append(('¿Trabajador nocturno según el ET?: %s' % quien, H_NOC, 'G%d' % (N_INI + i), 'salida'))
        m.append(('Horas de 22:00 a 6:00 a la semana: %s' % quien, H_NOC, 'C%d' % (N_INI + i), 'salida'))
        m.append(('Plus de nocturnidad al año: %s' % quien, H_NOC, 'L%d' % (N_INI + i), 'salida'))
    return m


# ==========================================================================
# Cierre: guardar, cachear, verificar, cruces, lista negra y mapa
# ==========================================================================
_PROHIBIDAS = ('INDIRECT', 'COUNTA', 'PMT(', 'OFFSET', 'XLOOKUP', 'LET(', 'LAMBDA', 'RANK(',
               'NETWORKDAYS', 'IRR(')
RX_CONST_DEC = re.compile(r'[*/]\s*\d{1,3}[.,]\d{1,4}(?!\d)')
#: Enteros que pueden ir en una fórmula: el cero de las comparaciones y de MAX(0;...) y el
#: uno de «uno más» (1 + Seguridad Social). Cualquier otro número tecleado aborta.
_ENTEROS_OK = {'0', '1'}
_RX_NUM = re.compile(r'(?<![A-Z$0-9.])(\d+(?:\.\d+)?)(?![0-9]*[(!])')


def _textos(wbf):
    for ws in wbf.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if c.__class__.__name__ == 'MergedCell':
                    continue
                if isinstance(c.value, str):
                    yield '%s!%s' % (ws.title, c.coordinate), c.value
                if c.comment is not None:
                    yield '%s!%s (nota)' % (ws.title, c.coordinate), c.comment.text


def _gate_lista_negra(wbf):
    fallos = []
    rotulos = set(D.rotulo_cruce(c) for c in CRUCES_IN)
    for donde, txt in _textos(wbf):
        bajo = txt.lower()
        for aguja, motivo in D.LISTA_NEGRA:
            if aguja.startswith('re:'):
                if re.search(aguja[3:], bajo):
                    fallos.append('%s: /%s/ (%s)' % (donde, aguja[3:], motivo))
            elif aguja.lower() in bajo:
                fallos.append('%s: «%s» (%s)' % (donde, aguja, motivo))
        for pid in RX_ID.findall(txt):
            if pid in D.IDS_PROHIBIDOS:
                fallos.append('%s: id prohibido %s' % (donde, pid))
        if 'trae aquí la cifra de' in bajo:
            fallos.append('%s: rótulo de cruce con la fórmula de la hermana' % donde)
        if 'cópialo del libro' in bajo and txt.strip() not in rotulos:
            fallos.append('%s: rótulo de cruce que no está en CRUCES: %s' % (donde, txt[:80]))
        if 'churrero' in bajo and 'categor' in bajo:
            fallos.append('%s: «categoría» junto a churrero (D9)' % donde)
    if fallos:
        raise SystemExit('LISTA NEGRA / IDS PROHIBIDOS en %s:\n  %s' % (NOMBRE, '\n  '.join(fallos[:40])))


def _gate_cruces(wbf, wbv):
    if D._ERROR_CRUCES:
        raise SystemExit('datos_ejemplo no pudo calcular los cruces: %s' % D._ERROR_CRUCES)
    fallos, filas = [], []
    if len(CRUCE_VERDES) != len(CRUCES_IN) or len(CRUCE_CUADRES) != len(CRUCES_IN):
        fallos.append('%d cruces, %d verdes de cruce y %d filas de CUADRE'
                      % (len(CRUCES_IN), len(CRUCE_VERDES), len(CRUCE_CUADRES)))
    for c in CRUCES_IN:
        if c['hoja_receptor'] != H_PAR:
            fallos.append('%s: la hoja receptora es %s y no «Parámetros»' % (c['x'], c['hoja_receptor']))
        fila = CRUCE_ROW[c['calcula']]
        cel_v = wbf[H_PAR]['B%d' % fila]
        if not motor.es_verde(cel_v) or cel_v.protection.locked:
            fallos.append('%s: Parámetros!B%d no es verde y editable' % (c['x'], fila))
        if wbf[H_PAR]['D%d' % fila].value != D.rotulo_cruce(c):
            fallos.append('%s: Parámetros!D%d no lleva el rótulo %r' % (c['x'], fila, D.rotulo_cruce(c)))
        v = wbv[H_PAR]['B%d' % fila].value
        esperado = c['valor_defecto']
        if not isinstance(v, (int, float)) or abs(v - esperado) > max(1e-9, abs(esperado) * 1e-9):
            fallos.append('%s: Parámetros!B%d = %r y datos_ejemplo dice %r' % (c['x'], fila, v, esperado))
        fc = PR['cuadre_' + c['calcula']]
        rot = wbf[H_PAR]['A%d' % fc].value
        if not (isinstance(rot, str) and rot.startswith('CUADRE')):
            fallos.append('%s: Parámetros!A%d no es una fila de CUADRE' % (c['x'], fc))
        if wbv[H_PAR]['D%d' % fc].value != 'CUADRA':
            fallos.append('%s: con los datos de El Molinete el CUADRE de la fila %d dice %r'
                          % (c['x'], fc, wbv[H_PAR]['D%d' % fc].value))
        if not motor.es_verde(wbf[H_PAR]['B%d' % fc]):
            fallos.append('%s: la celda «lo que publica hoy el origen» de la fila %d no es verde' % (c['x'], fc))
    for _h, coord in CONTRASTES:
        if wbv[H_PAR][coord].value != 'COHERENTE':
            fallos.append('el contraste %s dice %r con los datos de El Molinete'
                          % (coord, wbv[H_PAR][coord].value))
    for c in CRUCES_OUT:
        hoja, celda = c['hoja_origen'], c['celda_origen']
        rot = wbf[hoja]['K' + celda[1:]].value
        if not isinstance(rot, str) or not rot.strip():
            fallos.append('%s: %s!K%s sin rótulo' % (c['x'], hoja, celda[1:]))
        v = wbv[hoja][celda].value
        esperado = c['valor_defecto']
        if not isinstance(v, (int, float)) or abs(v - esperado) > max(1e-6, abs(esperado) * 1e-9):
            fallos.append('%s: %s!%s = %r y datos_ejemplo dice %r' % (c['x'], hoja, celda, v, esperado))
        filas.append((c['x'], hoja, celda, v, esperado))
    for ws in wbf.worksheets:
        if ws.title == H_PAR:
            continue
        for row in ws.iter_rows():
            for cel in row:
                if isinstance(cel.value, str) and 'cópialo del libro' in cel.value:
                    fallos.append('%s!%s: rótulo de cruce fuera de «Parámetros»' % (ws.title, cel.coordinate))
    if fallos:
        raise SystemExit('CRUCES ROTOS en %s:\n  %s' % (NOMBRE, '\n  '.join(fallos)))
    return filas


def _gate_contra_datos(wbv):
    """Las salidas principales, al céntimo contra `datos_ejemplo`."""
    pares = [
        ('coste anual de plantilla', wbv[H_COS]['B%d' % K_TOTAL].value, D.coste_plantilla_anual()),
        ('coste hora de la plantilla', wbv[H_COS]['B%d' % K_CHORA].value, D.coste_hora_mano_obra()),
        ('coste de los puestos', wbv[H_COS]['B%d' % K_PUESTOS].value, D.coste_puestos_anual()),
        ('horas de sustitución', wbv[H_COS]['B%d' % K_HSUS].value, D.horas_sustitucion_anio()),
        ('horas contratadas', wbv[H_COS]['B%d' % K_HORAS].value, D.horas_plantilla_anio()),
        ('horas del titular', wbv[H_TIT]['B%d' % T_HSEM].value, D.horas_titular()),
        ('madrugada del titular', wbv[H_TIT]['B%d' % T_MADRU].value, D.horas_titular_0_8()),
        ('huecos de cobertura', wbv[H_CUA]['B%d' % CR_HUE].value, float(len(D.huecos_de_cobertura()))),
    ]
    for i, p in enumerate(PLANTILLA):
        pid = p['id']
        a = D.analisis_nocturnidad(pid)
        pares += [
            ('%s horas semana' % pid, wbv[H_COS]['G%d' % (K_INI + i)].value, D.horas_semana(pid)),
            ('%s plus anual' % pid, wbv[H_NOC]['L%d' % (N_INI + i)].value, D.plus_nocturnidad_anual(pid)),
            ('%s salario anual' % pid, wbv[H_COS]['N%d' % (K_INI + i)].value, D.salario_anual(pid)),
            ('%s coste puesto' % pid, wbv[H_COS]['Q%d' % (K_INI + i)].value, D.coste_puesto_anual(pid)),
            ('%s horas 22-6' % pid, wbv[H_NOC]['C%d' % (N_INI + i)].value, a['horas_22_6_semana']),
            ('%s días con 3 h' % pid, wbv[H_NOC]['E%d' % (N_INI + i)].value, a['dias_con_3h_o_mas']),
            ('%s días trabajados' % pid, wbv[H_NOC]['D%d' % (N_INI + i)].value, a['dias_trabajados']),
            ('%s horas 22-0' % pid, wbv[H_NOC]['H%d' % (N_INI + i)].value, a['horas_plus_22_0']),
            ('%s horas 0-8' % pid, wbv[H_NOC]['I%d' % (N_INI + i)].value, a['horas_plus_0_8']),
            ('%s días jornada nocturna' % pid, wbv[H_NOC]['J%d' % (N_INI + i)].value,
             a['dias_jornada_nocturna_convenio']),
        ]
    fallos = ['%s: libro %r, datos_ejemplo %r' % (n, v, e) for n, v, e in pares
              if not isinstance(v, (int, float)) or abs(v - e) > max(1e-6, abs(e) * 1e-9)]
    for i, p in enumerate(PLANTILLA):
        v = wbv[H_NOC]['G%d' % (N_INI + i)].value
        e = D.analisis_nocturnidad(p['id'])['trabajador_nocturno_et']
        if p['es_titular']:
            e = VER_NA      # R-20: el Estatuto no se aplica al titular autónomo
        if v != e:
            fallos.append('%s veredicto ET: libro %r, datos_ejemplo %r' % (p['id'], v, e))
        pares.append(('%s veredicto ET' % p['id'], v, e))
    if fallos:
        raise SystemExit('SALIDAS que no cuadran con datos_ejemplo:\n  ' + '\n  '.join(fallos))
    return pares


def cerrar(wb, mapa_celdas_):
    wb.properties.creator = 'AI Chef Pro'
    wb.properties.lastModifiedBy = 'AI Chef Pro'
    wb.properties.title = TITULO
    wb.properties.subject = SUBJECT

    enmiendas_churreria.sanear_libro(wb)       # R-09: sin texto interno de la fábrica
    CC.barrer_cp1252(wb)
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and (motor.NARROW in c.value or motor.NOBRK in c.value):
                    raise SystemExit('%s!%s lleva espacio fino o guion no separable' % (ws.title, c.coordinate))
                if c.comment is not None:
                    try:
                        c.comment.text.encode('cp1252')
                    except UnicodeEncodeError:
                        raise SystemExit('%s!%s: la nota no cabe en WinAnsi' % (ws.title, c.coordinate))
    for ws in wb.worksheets:
        motor.retirar_verde_de_calculadas(ws)
        motor.proteger(ws)

    if wb.sheetnames[0] != H_INS:
        raise SystemExit('La primera hoja no es «Instrucciones»')
    if tuple(wb.sheetnames) != tuple(D.HOJAS[LIBRO]):
        raise SystemExit('Hojas %r frente a las de datos_ejemplo %r' % (wb.sheetnames, D.HOJAS[LIBRO]))

    destino = CC.BUILD_DIR
    if not os.path.isdir(destino):
        os.makedirs(destino)
    ruta = os.path.join(destino, NOMBRE + '.xlsx')

    otros = [os.path.splitext(f)[0] for n, f in D.LIBROS.items() if n != LIBRO]
    problemas = []
    for hoja, coord, form in motor.REGISTRO:
        arriba = form.upper()
        for p in _PROHIBIDAS:
            if p in arriba:
                problemas.append('%s!%s usa %s' % (hoja, coord, p))
        bajo = form.lower()
        if '.xlsx' in bajo or '[' in form or any(o in bajo for o in otros):
            problemas.append('%s!%s parece referencia a otro libro: %s' % (hoja, coord, form[:70]))
        sin_texto = re.sub(r'"[^"]*"', '', form)
        sin_refs = re.sub(r"'[^']*'!", '', sin_texto)
        for m in RX_CONST_DEC.finditer(sin_refs):
            problemas.append('%s!%s constante tecleada %s en %s' % (hoja, coord, m.group(0), form[:70]))
        for m in _RX_NUM.finditer(re.sub(r'\$?[A-Z]{1,2}\$?\d+', '', sin_refs)):
            if m.group(1) not in _ENTEROS_OK:
                problemas.append('%s!%s número tecleado %s en %s' % (hoja, coord, m.group(1), form[:70]))
        if '/' in sin_texto and 'IFERROR' not in arriba:
            problemas.append('%s!%s división sin IFERROR: %s' % (hoja, coord, form[:70]))
    if problemas:
        raise SystemExit('FÓRMULAS con problemas en %s:\n  %s' % (NOMBRE, '\n  '.join(problemas[:40])))
    wb.save(ruta)

    res = subprocess.run([sys.executable, os.path.join(RAIZ, 'inject_cache.py'), ruta],
                         stdout=subprocess.PIPE, stderr=subprocess.STDOUT)
    salida_cache = res.stdout.decode('utf-8', 'replace')
    if res.returncode != 0:
        raise SystemExit('inject_cache falló:\n' + salida_cache)

    wbf = openpyxl.load_workbook(ruta)
    wbv = openpyxl.load_workbook(ruta, data_only=True)
    sin_valor, con_error, sin_dato, pendientes = [], [], [], []
    for hoja, coord, form in motor.REGISTRO:
        v = wbv[hoja][coord].value
        if v is None:
            pendientes.append((hoja, coord, form))
        elif isinstance(v, str) and v.startswith('#'):
            con_error.append('%s!%s  %s -> %s' % (hoja, coord, form[:60], v))
    if pendientes:
        from pycel import ExcelCompiler
        xl = ExcelCompiler(ruta)
        for hoja, coord, form in pendientes:
            try:
                r = xl.evaluate("'%s'!%s" % (hoja, coord))
            except Exception as e:                            # noqa: BLE001
                sin_valor.append('%s!%s  %s -> pycel: %s' % (hoja, coord, form[:60], e))
                continue
            if r == '' or (r is None and '""' in form):
                sin_dato.append('%s!%s' % (hoja, coord))
            else:
                sin_valor.append('%s!%s  %s -> %r' % (hoja, coord, form[:60], r))
    if sin_valor or con_error:
        raise SystemExit('VERIFICACIÓN data_only FALLIDA en %s\n  sin valor (%d):\n   %s\n  con error '
                         '(%d):\n   %s' % (NOMBRE, len(sin_valor), '\n   '.join(sin_valor[:40]),
                                           len(con_error), '\n   '.join(con_error[:40])))
    incoherentes = CC.gate_sum_rango(wbv)
    if incoherentes:
        raise SystemExit('TOTALES con caché incoherente:\n  ' + '\n  '.join(incoherentes))

    verdes_vacias, n_verdes = [], 0
    for ws in wbf.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if c.__class__.__name__ == 'MergedCell':
                    continue
                if motor.es_verde(c):
                    n_verdes += 1
                    if c.protection.locked:
                        verdes_vacias.append('%s!%s bloqueada' % (ws.title, c.coordinate))
                    if c.value is None or (isinstance(c.value, str) and not c.value.strip()):
                        verdes_vacias.append('%s!%s' % (ws.title, c.coordinate))
                    if c.data_type == 'f':
                        verdes_vacias.append('%s!%s verde con fórmula' % (ws.title, c.coordinate))
    if verdes_vacias:
        raise SystemExit('CELDAS VERDES VACÍAS, BLOQUEADAS O CON FÓRMULA:\n  ' + ', '.join(verdes_vacias))

    sin_dv = []
    for ws in wbf.worksheets:
        cubiertas = set()
        for dv in ws.data_validations.dataValidation:
            if not dv.showErrorMessage:
                sin_dv.append('%s: DV sin showErrorMessage (%s)' % (ws.title, dv.sqref))
            if dv.type == 'list' and dv.formula1 and dv.formula1.startswith('"'):
                sin_dv.append('%s: DV de lista con comas (%s)' % (ws.title, dv.sqref))
            for rng in str(dv.sqref).split():
                for fila in ws[rng] if ':' in rng else [[ws[rng]]]:
                    for cel in fila:
                        cubiertas.add(cel.coordinate)
        for row in ws.iter_rows():
            for c in row:
                if c.__class__.__name__ == 'MergedCell':
                    continue
                lista = ((ws.title == H_CUA and c.column_letter in ('A', 'B', 'E')
                          and CT_INI <= c.row <= CT_FIN)
                         or (ws.title == H_COS and c.column_letter in ('D', 'E', 'F')
                             and K_INI <= c.row <= K_FIN)
                         or (ws.title == H_PAR and c.row == PR['cubrir']))
                if (c.__class__.__name__ != 'MergedCell' and motor.es_verde(c)
                        and c.coordinate not in cubiertas
                        and (isinstance(c.value, (int, float)) or lista)):
                    sin_dv.append('%s!%s verde sin validación' % (ws.title, c.coordinate))
    if sin_dv:
        raise SystemExit('VALIDACIÓN DE DATOS:\n  ' + '\n  '.join(sin_dv[:40]))

    _gate_lista_negra(wbf)
    contrato = _gate_cruces(wbf, wbv)
    contra_datos = _gate_contra_datos(wbv)

    requeridos = [i for i, (_t, libros) in D.IDS_LEGALES_REQUERIDOS.items() if LIBRO in libros]
    puestos = set(pid for _h, _c, pid in NOTAS_LEGALES)
    faltan = [i for i in requeridos if i not in puestos]
    if faltan:
        raise SystemExit('Faltan notas legales del libro 8: %s' % ', '.join(faltan))

    mapa = {}
    for etiqueta, hoja, coord, tipo in mapa_celdas_:
        if tipo not in ('entrada', 'salida', 'parametro'):
            raise SystemExit('Tipo de mapa no válido: %r' % tipo)
        v = wbv[hoja][coord].value
        if v is None:
            raise SystemExit('El mapa cita %s!%s («%s») y está VACÍA' % (hoja, coord, etiqueta))
        if etiqueta in mapa:
            raise SystemExit('Etiqueta de mapa repetida: %r' % etiqueta)
        mapa[etiqueta] = {'ref': '%s!%s!%s' % (FICHERO, hoja, coord), 'valor': v, 'tipo': tipo}
    if len(mapa) < 25:
        raise SystemExit('El mapa tiene %d etiquetas y el mínimo es 25' % len(mapa))
    for x, hoja, celda, _v, _e in contrato:
        if not any(d['ref'] == '%s!%s!%s' % (FICHERO, hoja, celda) for d in mapa.values()):
            raise SystemExit('El mapa no cita la celda de origen %s!%s (%s)' % (hoja, celda, x))
    with open(os.path.join(destino, 'mapa-' + NOMBRE + '.json'), 'w', encoding='utf-8') as fh:
        fh.write(json.dumps(mapa, ensure_ascii=False, indent=1))

    return {'ruta': ruta, 'hojas': len(wbf.worksheets), 'formulas': len(motor.REGISTRO),
            'sin_dato': len(sin_dato), 'verdes': n_verdes, 'verdes_vacias': verdes_vacias,
            'notas_legales': len(NOTAS_LEGALES), 'mapa': len(mapa), 'contrato': contrato,
            'contra_datos': contra_datos, 'cruces_in': len(CRUCE_VERDES),
            'cuadres': len(CRUCE_CUADRES),
            'cache': (salida_cache.strip().splitlines() or [''])[-1]}


# ==========================================================================
# Demostraciones con pycel
# ==========================================================================
def demo(ruta):
    from pycel import ExcelCompiler
    ok, fallos = [], []

    def prueba(nombre, cond, detalle=''):
        (ok if cond else fallos).append(nombre + (' - ' + detalle if detalle else ''))

    def ev(xl, hoja, coord):
        return xl.evaluate("'%s'!%s" % (hoja, coord))

    def pon(xl, hoja, coord, v):
        xl.set_value("'%s'!%s" % (hoja, coord), v)

    i_p3 = [p['id'] for p in PLANTILLA].index('P3')
    i_p5 = [p['id'] for p in PLANTILLA].index('P5')
    f_p3_sab = CT_INI + TURNOS.index(('P3', 'S', 4.5, 12.5, 'L'))
    f_p3_dom = CT_INI + TURNOS.index(('P3', 'D', 4.5, 12.5, 'L'))

    # 1. Adelantar el turno del fin de semana del churrero/a 2 de las 4:30 a las 3:00: pasa a
    #    tener tres horas en el periodo del ET esos dos días -> deja de ser «no», cobra más
    #    plus y el coste de plantilla (X15) sube.
    xl = ExcelCompiler(ruta)
    x0 = ev(xl, H_COS, 'L5')
    g0 = ev(xl, H_NOC, 'G%d' % (N_INI + i_p3))
    l0 = ev(xl, H_NOC, 'L%d' % (N_INI + i_p3))
    pon(xl, H_CUA, 'C%d' % f_p3_sab, 3.0)
    pon(xl, H_CUA, 'C%d' % f_p3_dom, 3.0)
    g1 = ev(xl, H_NOC, 'G%d' % (N_INI + i_p3))
    l1 = ev(xl, H_NOC, 'L%d' % (N_INI + i_p3))
    x1 = ev(xl, H_COS, 'L5')
    prueba('X15 = datos_ejemplo; con el turno del fin de semana a las 3:00 el churrero/a 2 pasa '
           'de «no» a otra figura, cobra más plus y sube el coste de plantilla',
           abs(x0 - D.coste_plantilla_anual()) < 1e-6 and g0 == VER_NO and g1 != VER_NO
           and l1 > l0 and x1 > x0,
           '%r -> %r · plus %.2f -> %.2f · X15 %.2f -> %.2f' % (g0, g1, l0, l1, x0, x1))

    # 2. Un coste hora copiado del libro 3 un 20 % más alto: el CUADRE de X8 y el contraste H se
    #    ponen en rojo; el CUADRE de las personas en la línea sigue en verde.
    xl = ExcelCompiler(ruta)
    fh = PR['cuadre_coste_hora_mano_obra']
    ff = PR['cuadre__personas_pico']
    ch = PR['contraste_coste_hora_mano_obra']
    h0, f0 = ev(xl, H_PAR, 'D%d' % fh), ev(xl, H_PAR, 'D%d' % ff)
    k0 = ev(xl, H_PAR, 'B%d' % ch)
    pon(xl, H_PAR, 'B%d' % CRUCE_ROW['coste_hora_mano_obra'], D.coste_hora_mano_obra() * 1.2)
    h1, f1 = ev(xl, H_PAR, 'D%d' % fh), ev(xl, H_PAR, 'D%d' % ff)
    k1 = ev(xl, H_PAR, 'B%d' % ch)
    prueba('un coste hora del libro 3 un 20 % por encima pone en rojo el CUADRE de X8 y el contraste H '
           'y deja en verde el de las personas en la línea',
           h0 == 'CUADRA' and h1.startswith('REVISA') and f0 == 'CUADRA' and f1 == 'CUADRA'
           and k0 == 'COHERENTE' and k1 == KO['H'], '%r -> %r' % (h0, h1[:12]))

    # 2 bis. R-01: un refuerzo copiado diez veces mayor (560 h en vez de 56) sube el coste de plantilla
    #        (X15) en silencio hacia el libro 6; antes el CUADRE decía «CUADRA». Ahora avisa.
    xl = ExcelCompiler(ruta)
    fr = PR['cuadre_horas_refuerzo_pico']
    r0 = ev(xl, H_PAR, 'D%d' % fr)
    x0 = ev(xl, H_COS, 'L5')
    pon(xl, H_PAR, 'B%d' % CRUCE_ROW['horas_refuerzo_pico'], D.horas_refuerzo_pico() * 10)
    r1 = ev(xl, H_PAR, 'D%d' % fr)
    x1 = ev(xl, H_COS, 'L5')
    prueba('560 horas de refuerzo copiadas en vez de 56: el coste de plantilla (X15) sube Y el CUADRE '
           'del refuerzo dice REVISA', r0 == 'CUADRA' and r1.startswith('REVISA') and x1 > x0 + 6000,
           'X15 %.2f -> %.2f · %r -> %r' % (x0, x1, r0, r1[:12]))

    # 3. El salario base del nivel V a 500 EUR: el refuerzo queda por debajo del SMI, el
    #    semáforo lo dice y su coste se queda en el suelo (SMI a media jornada más la SS).
    xl = ExcelCompiler(ruta)
    s0 = ev(xl, H_COS, 'P%d' % (K_INI + i_p5))
    ev(xl, H_COS, 'Q%d' % (K_INI + i_p5))
    pon(xl, H_PAR, 'B%d' % NIV_ROWS['V'], 500.0)
    s = ev(xl, H_COS, 'P%d' % (K_INI + i_p5))
    q = ev(xl, H_COS, 'Q%d' % (K_INI + i_p5))
    suelo = D.dato(D.CONVENIO, 'smi_anual') * 0.5 * (1 + D.P('ss_empresa'))
    prueba('nivel V a 500 EUR/mes: el refuerzo cae por debajo del SMI y su coste es el suelo',
           s0 == SMI_OK and s == SMI_BAJO and abs(q - suelo) < 1e-6,
           '%r -> %r · coste %.2f (suelo %.2f)' % (s0, s, q, suelo))

    # 4. Quien cierra el viernes deja la freidora a las 21:00 (fin del turno de noche a las
    #    21): aparecen huecos de cobertura (ocho medias horas de la noche del viernes).
    xl = ExcelCompiler(ruta)
    f_noche = CT_INI + TURNOS.index(('P4', 'V', 21.0, 25.0, 'L'))
    u0 = ev(xl, H_CUA, 'B%d' % CR_HUE)
    v0 = ev(xl, H_CUA, 'B%d' % CR_VER)
    pon(xl, H_CUA, 'D%d' % f_noche, 21.0)
    u1 = ev(xl, H_CUA, 'B%d' % CR_HUE)
    v1 = ev(xl, H_CUA, 'B%d' % CR_VER)
    prueba('sin el turno de noche del viernes aparecen ocho medias horas sin cubrir',
           u0 == 0 and u1 == 8 and v0.startswith('Sin huecos') and v1.startswith('HAY HUECOS'),
           '%r -> %r' % (u0, u1))

    # 5. Copiar del libro 1 tres personas en la línea del pico: el mínimo del pico sube, salen
    #    huecos en el desayuno del fin de semana y el CUADRE de X6 se pone en rojo.
    xl = ExcelCompiler(ruta)
    fe = PR['cuadre__personas_pico']
    ce = PR['contraste__personas_pico']
    e0 = ev(xl, H_PAR, 'D%d' % fe)
    ce0 = ev(xl, H_PAR, 'B%d' % ce)
    ev(xl, H_CUA, 'B%d' % CR_HUE)
    pon(xl, H_PAR, 'B%d' % CRUCE_ROW['_personas_pico'], 3)
    e1 = ev(xl, H_PAR, 'D%d' % fe)
    ce1 = ev(xl, H_PAR, 'B%d' % ce)
    u2 = ev(xl, H_CUA, 'B%d' % CR_HUE)
    prueba('tres personas en la línea del pico: CUADRE de X6 en rojo, contraste E incoherente y 16 '
           'medias horas sin cubrir (de 6:00 a 10:00, sábado y domingo)',
           e0 == 'CUADRA' and e1.startswith('REVISA') and ce0 == 'COHERENTE' and ce1 == KO['E']
           and u2 == 16, 'huecos %r' % u2)
    return ok, fallos


# ==========================================================================
def main():
    requeridos = [i for i, (_t, libros) in D.IDS_LEGALES_REQUERIDOS.items() if LIBRO in libros]
    D.gate_legal(requeridos)
    for c in CRUCES_IN + CRUCES_OUT:
        if (c['origen'], c['receptor']) not in D.ARISTAS_SPEC:
            raise SystemExit('%s: arista %s no está en la SPEC' % (c['x'], (c['origen'], c['receptor'])))
    if D.huecos_de_cobertura() or D.problemas_del_cuadrante():
        raise SystemExit('El cuadrante de datos_ejemplo tiene huecos o problemas')
    wb = Workbook()
    wb.remove(wb.active)
    hoja_instrucciones(wb)
    hoja_parametros(wb)
    hoja_cuadrante(wb)
    hoja_nocturnidad(wb)
    hoja_coste(wb)
    hoja_identifica(wb)
    hoja_titular(wb)
    orden = list(D.HOJAS[LIBRO])
    wb._sheets.sort(key=lambda ws: orden.index(ws.title))

    res = cerrar(wb, mapa_celdas())
    print('escrito: %s' % res['ruta'])
    print('hojas: %d · fórmulas: %d · celdas verdes: %d · verdes vacías: %d'
          % (res['hojas'], res['formulas'], res['verdes'], len(res['verdes_vacias'])))
    print('fórmulas que devuelven «sin dato» a propósito: %d' % res['sin_dato'])
    print('notas legales: %d · etiquetas en el mapa: %d' % (res['notas_legales'], res['mapa']))
    print('cruces recibidos: %d (con %d filas de CUADRE) · cruces de origen: %d'
          % (res['cruces_in'], res['cuadres'], len(res['contrato'])))
    print('inject_cache: %s' % res['cache'])
    for x, hoja, celda, v, e in res['contrato']:
        print('  origen %-5s %s!%s = %.6f (datos_ejemplo %.6f)' % (x, hoja, celda, v, e))
    for n, v, e in res['contra_datos'][:8]:
        print('  contra datos_ejemplo · %s = %r' % (n, v))
    ok, fallos = demo(res['ruta'])
    for t in ok:
        print('  demo OK    · %s' % t)
    for t in fallos:
        print('  demo FALLA · %s' % t)
    if fallos:
        raise SystemExit('demos con fallos: %d' % len(fallos))
    return res, len(ok), len(fallos)


if __name__ == '__main__':
    main()
