#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
gen_checklist-legal-fritura-y-licencias.py — libro 7 de «Cómo Montar una
Churrería-Chocolatería» (SPEC §2.1, fila 7; molde H+N de §2.3).

Hojas (las de `datos_ejemplo.HOJAS[7]`, en su orden): Instrucciones · Checklist
Legal (F1-F6) · Árbol IAE y CNAE por Formato · Régimen de Apertura y Horario ·
Licencia, Humos y Gas · Árbol de Registro Sanitario · Ferias y Venta Ambulante ·
Alérgenos y Aceite Compartido · PRL de la Fritura · Registro de Formación ·
Cronograma y Ruta Crítica.

QUÉ DECIDE ESTE LIBRO
---------------------
«¿Qué papel me toca y en qué orden?». Se calca de
`guia-chocolateria/gen_checklist-legal-licencias-y-cacao.py` (checklist por
fases, árbol de registro, registro de formación y cronograma con ruta crítica) y
de `_comun_libros_8_9.py` (cierre). FUERA, por la SPEC: cadmio, EUDR y ruta
doméstica. NUEVO: árbol IAE y CNAE, régimen y horario, licencia/humos/gas,
ferias, alérgenos y aceite compartido, PRL de la fritura.

LAS REGLAS DURAS QUE GOBIERNAN ESTE LIBRO
-----------------------------------------
1. D13 / A4: la PRIMERA fila del checklist es «¿la extracción necesita
   proyecto?» (licencia de obra y técnico) y la segunda «¿tu actividad está en
   el Anexo de la Ley 12/2012?». La actividad del despacho va por declaración
   responsable; las obras con proyecto, NO.
2. V-03 rama B: el CNAE va SIEMPRE como pregunta al asesor, con doble código
   (CNAE-2025 + CNAE-2009) y la trampa del 47.81 (`CUN-15`). Ninguna fórmula
   devuelve «tu CNAE es».
3. A3 / D14: ninguna cuota del IAE y nada de la nota de los locales de
   temporada del 676 en este libro. La exención (`CHN-73`) sí: alta sí, pago
   casi nunca.
4. D10: el horario de Madrid es un EJEMPLO, sin extrapolar, y este libro NO pide
   el horario: remite a `temporada-franjas-y-ferias.xlsx!Franjas del Día`.
5. A6: la bombona de menos de 15 kg sólo para un aparato «de utilización móvil»,
   y sólo en la caseta o feria. En el local, instalación receptora e inspección
   periódica (`CHN-48`).
6. D12 / A14: la acrilamida del churro es buena práctica; la parte A del anexo
   II si fríes patatas; la parte B sólo con las condiciones del art. 2.3.
7. Ferias con la Ley 7/1996 (`CUN-20`); el RD de venta ambulante de 2010 está
   derogado (`CUN-19`) y no se cita ninguno de sus artículos.
8. Registro de formación, nunca la palabra prohibida que empieza por «c» y que
   la familia tiene en su lista negra (`CHN-69`).
9. El plazo de entrega de la maquinaria crítica vive SÓLO aquí (fuente única del
   paquete): una celda verde en «Cronograma y Ruta Crítica» de la que beben el
   hito H7 y la fila del checklist que lo pide.
10. El libro 7 no depende de ningún otro (D16): CERO celdas de cruce, y un gate
   que aborta si alguna celda dice «cópialo del libro» o «trae aquí la cifra».

Salida: build/checklist-legal-fritura-y-licencias.xlsx
        + build/mapa-checklist-legal-fritura-y-licencias.json
Uso:    /usr/local/bin/python3 gen_checklist-legal-fritura-y-licencias.py
Via: Claude Code
"""
import datetime
import importlib.util
import json
import os
import re
import subprocess
import sys

import openpyxl
from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill

AQUI = os.path.dirname(os.path.abspath(__file__))
RAIZ = os.path.normpath(os.path.join(AQUI, '..'))
for _p in (os.path.join(RAIZ, 'guias-v2_0'), AQUI):
    if _p not in sys.path:
        sys.path.insert(0, _p)

import _comun_churreria as C                                   # noqa: E402
import motor                                                   # noqa: E402

_spec = importlib.util.spec_from_file_location(
    'datos_ejemplo_churreria', os.path.join(AQUI, 'datos_ejemplo.py'))
D = importlib.util.module_from_spec(_spec)
_spec.loader.exec_module(D)
import enmiendas_churreria                                     # noqa: E402
enmiendas_churreria.aplicar(D)      # contrato de cruces enmendado y notas legales sin lo vetado

motor.CTX['producto'] = C.PID

LIBRO = 7
NOMBRE = D.LIBROS[LIBRO][:-len('.xlsx')]
TITULO = 'Checklist Legal, Fritura y Licencias'
SUBJECT = C.PRODUCTO + ' · Versión 1.0 · octubre 2026'

(H_INS, H_CHK, H_IAE, H_REG, H_HUM, H_SAN, H_FER, H_ALE, H_PRL, H_FOR,
 H_GAN) = D.HOJAS[LIBRO]

HECHO, PENDIENTE, NOAPLICA = 'Hecho', 'Pendiente', 'N/A'
SI, NO, NOSE = 'Sí', 'No', 'No lo sé'
REVISA = 'Revisa la etiqueta'
SIN_IMPORTE = 'A presupuestar'
SIN_PAGAR = 'Sin pagar'
SIN_DEP = '-'
FECHA_BASE = datetime.date(2026, 10, 3)
FMT_HORA = 'hh:mm'

CCAA = ['Andalucía', 'Aragón', 'Principado de Asturias', 'Illes Balears',
        'Canarias', 'Cantabria', 'Castilla-La Mancha', 'Castilla y León',
        'Cataluña', 'Comunitat Valenciana', 'Extremadura', 'Galicia',
        'Comunidad de Madrid', 'Región de Murcia', 'Comunidad Foral de Navarra',
        'País Vasco', 'La Rioja', 'Ceuta', 'Melilla']
CCAA_EJEMPLO = 'Comunidad de Madrid'

#: Los ids legales que este libro cita según `datos_ejemplo` (los que llevan el 7).
IDS_LEGALES = [k for k, v in D.IDS_LEGALES_REQUERIDOS.items() if LIBRO in v[1]]
RX_NOTA = re.compile(r'\b(?:CUN|CHN)-(?:[MVD])?[0-9]+[a-z]*\b')


def ie(expr):
    return motor.iferror(expr)


def q(hoja):
    """Referencia a otra hoja del MISMO libro (nunca a otro fichero)."""
    return "'%s'!" % hoja


# ==========================================================================
# Cierre y entradas (calco de guia-chocolateria/_comun_libros_8_9.py)
# ==========================================================================
NOTAS_LEGALES = []
VERDES = []


def entrada(ws, coord, valor, fmt=None, etiqueta='', bold=None, align=None,
            wrap=None):
    """Celda VERDE de entrada, con su valor por defecto (ninguna verde vacía)."""
    if valor is None or (isinstance(valor, str) and not valor.strip()):
        raise SystemExit('entrada(%s!%s) sin valor por defecto' % (ws.title, coord))
    cel = motor.val(ws, coord, valor, fmt=fmt, verde_=True, bold=bold,
                    align=align, wrap=wrap)
    VERDES.append((ws.title, coord, etiqueta, valor))
    return cel


def _winansi(texto):
    out = []
    for ch in texto:
        try:
            ch.encode('cp1252')
        except UnicodeEncodeError:
            continue
        out.append(ch)
    return re.sub(r'\s+', ' ', ''.join(out)).strip()


def nota_celda(ws, coord, pid):
    """Nota LEGAL «Verificado el <fecha> · norma y artículo · URL» de
    `datos_ejemplo.nota_legal(id)` como comentario de la celda."""
    texto = D.nota_legal(pid)
    if not texto:
        raise SystemExit('nota_legal(%r) vacía: gate_legal debería haber '
                         'abortado antes' % pid)
    ws[coord].comment = Comment(_winansi(texto), 'AI Chef Pro', height=130,
                                width=470)
    NOTAS_LEGALES.append((ws.title, coord, pid))
    return texto


def notas_de_fuente(ws, coords, fuente):
    """Pone, celda a celda, las notas legales de los ids con nota de `fuente`."""
    ids = [i for i in RX_NOTA.findall(fuente or '') if D.nota_legal(i)]
    for coord, pid in zip(coords, ids):
        nota_celda(ws, coord, pid)
    return ids


def total_fila(ws, fila, letras, fondo=None):
    for letra in letras:
        c = ws[letra + str(fila)]
        c.font = Font(bold=True, size=c.font.size)
        c.fill = PatternFill('solid', fgColor=fondo or C.CREMA)


def barra_seccion(ws, fila, texto, letras):
    C.seccion(ws, 'A%d' % fila, texto)
    for letra in letras:
        motor.val(ws, letra + str(fila), '')
        ws[letra + str(fila)].fill = PatternFill('solid', fgColor=C.CABECERA)


_RX_SUM = re.compile(r'^=SUM\(([A-Z]{1,2})(\d+):([A-Z]{1,2})(\d+)\)$')
PROHIBIDAS = ('INDIRECT', 'COUNTA', 'PMT(', 'OFFSET', 'XLOOKUP', 'LET(',
              'LAMBDA', 'RANK(', 'NETWORKDAYS', 'IRR(')
RX_CRUCE = re.compile(r'cópialo del libro|trae aquí la cifra', re.I)


def cerrar(wb, mapa_celdas):
    """Guarda, inyecta caché, verifica `data_only`, pasa los gates y escribe el
    mapa. Devuelve el resumen como dict."""
    if wb.sheetnames != list(D.HOJAS[LIBRO]):
        raise SystemExit('Hojas fuera de orden: %r' % wb.sheetnames)
    wb.properties.creator = 'AI Chef Pro'
    wb.properties.lastModifiedBy = 'AI Chef Pro'
    wb.properties.title = TITULO
    wb.properties.subject = SUBJECT
    enmiendas_churreria.sanear_libro(wb)       # R-09: sin texto interno de la fábrica
    C.barrer_cp1252(wb)

    # --- gate: libro aislado (D16) y lista negra sobre el texto del libro ----
    textos = []
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str):
                    textos.append(('%s!%s' % (ws.title, c.coordinate), c.value))
    cruces = [t for t in textos if RX_CRUCE.search(t[1])]
    if cruces:
        raise SystemExit('El libro 7 no recibe cruces (D16) y dice: %r' % cruces[:5])
    negras = []
    for donde, t in textos:
        bajo = t.lower()
        for aguja, motivo in D.LISTA_NEGRA:
            hit = (re.search(aguja[3:], bajo) if aguja.startswith('re:')
                   else (aguja in bajo))
            if hit:
                negras.append('%s: %r (%s)' % (donde, aguja, motivo))
        for pid in D.IDS_PROHIBIDOS:
            if re.search(r'\b%s\b(?![a-z])' % re.escape(pid), t):
                negras.append('%s: id prohibido %s' % (donde, pid))
        if 'carnet' in bajo or 'carné' in bajo:
            negras.append('%s: la palabra del registro de formación' % donde)
        for pid in D.RX_ID.findall(t):
            if not D.id_existe_en_json_comun(pid):
                negras.append('%s: id inexistente %s' % (donde, pid))
    if negras:
        raise SystemExit('LISTA NEGRA en el libro 7:\n  ' + '\n  '.join(negras[:30]))

    # --- gate: cada id legal del libro 7 lleva al menos una nota ------------
    puestos = set(n[2] for n in NOTAS_LEGALES)
    faltan = [i for i in IDS_LEGALES if i not in puestos]
    if faltan:
        raise SystemExit('Ids legales del libro 7 sin ninguna nota en una celda: %s'
                         % faltan)

    for ws in wb.worksheets:
        motor.retirar_verde_de_calculadas(ws)
        motor.proteger(ws)

    os.makedirs(C.BUILD_DIR, exist_ok=True)
    ruta = os.path.join(C.BUILD_DIR, NOMBRE + '.xlsx')
    wb.save(ruta)

    res = subprocess.run([sys.executable, os.path.join(RAIZ, 'inject_cache.py'),
                          ruta], capture_output=True, text=True)
    salida_cache = (res.stdout or '') + (res.stderr or '')
    if res.returncode != 0:
        raise SystemExit('inject_cache falló:\n' + salida_cache)

    # --- verificación data_only ---------------------------------------------
    wbf = openpyxl.load_workbook(ruta)
    wbv = openpyxl.load_workbook(ruta, data_only=True)
    sin_valor, con_error, sin_dato, pendientes = [], [], [], []
    for hoja, coord, formula in motor.REGISTRO:
        v = wbv[hoja][coord].value
        if v is None:
            pendientes.append((hoja, coord, formula))
        elif isinstance(v, str) and v.startswith('#'):
            con_error.append('%s!%s  %s -> %s' % (hoja, coord, formula[:70], v))
    if pendientes:
        from pycel import ExcelCompiler
        xl = ExcelCompiler(ruta)
        for hoja, coord, formula in pendientes:
            try:
                r = xl.evaluate("'%s'!%s" % (hoja, coord))
            except Exception as e:                            # noqa: BLE001
                sin_valor.append('%s!%s  %s -> pycel: %s' % (hoja, coord, formula[:70], e))
                continue
            if r == '' or (r is None and formula.count('""')):
                sin_dato.append('%s!%s' % (hoja, coord))
            else:
                sin_valor.append('%s!%s  %s -> %r' % (hoja, coord, formula[:70], r))
    if sin_valor or con_error:
        raise SystemExit('VERIFICACIÓN data_only FALLIDA\n  sin valor (%d):\n   %s\n'
                         '  con error (%d):\n   %s'
                         % (len(sin_valor), '\n   '.join(sin_valor[:40]),
                            len(con_error), '\n   '.join(con_error[:40])))

    incoherentes = []
    for hoja, coord, formula in motor.REGISTRO:
        m = _RX_SUM.match(formula)
        if not m or m.group(1) != m.group(3):
            continue
        col, a, b = m.group(1), int(m.group(2)), int(m.group(4))
        esperado = sum(v for v in (wbv[hoja]['%s%d' % (col, r)].value
                                   for r in range(a, b + 1)) if isinstance(v, (int, float)))
        real = wbv[hoja][coord].value
        if not isinstance(real, (int, float)) or abs(real - esperado) > 0.01:
            incoherentes.append('%s!%s' % (hoja, coord))
    if incoherentes:
        raise SystemExit('TOTALES con caché incoherente: %s' % incoherentes)

    verdes_vacias, n_verdes = [], 0
    for ws in wbf.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if c.__class__.__name__ == 'MergedCell':
                    continue
                if motor.es_verde(c):
                    n_verdes += 1
                    if c.value is None or (isinstance(c.value, str) and not c.value.strip()):
                        verdes_vacias.append('%s!%s' % (ws.title, c.coordinate))
    if verdes_vacias:
        raise SystemExit('CELDAS VERDES VACÍAS: %s' % verdes_vacias[:40])

    usos = []
    for hoja, coord, formula in motor.REGISTRO:
        arriba = formula.upper()
        for p in PROHIBIDAS:
            if p in arriba:
                usos.append('%s!%s usa %s' % (hoja, coord, p.rstrip('(')))
        if '.XLSX' in arriba or '[' in formula:
            usos.append('%s!%s NOMBRA OTRO FICHERO' % (hoja, coord))
        sin_txt = re.sub(r'"[^"]*"', '', formula)
        for lit in motor.RX_LITERAL.findall(sin_txt):
            if lit not in motor.LITERALES_TOLERADOS:
                usos.append('%s!%s constante %s dentro de la fórmula' % (hoja, coord, lit))
    if usos:
        raise SystemExit('Fórmulas PROHIBIDAS, referencia externa o constante:\n  '
                         + '\n  '.join(usos))

    # --- mapa -----------------------------------------------------------------
    mapa = {}
    for etiqueta, hoja, coord, tipo in mapa_celdas:
        if tipo not in ('entrada', 'salida', 'parametro'):
            raise SystemExit('Tipo de mapa no válido: %r' % tipo)
        v = wbv[hoja][coord].value
        if v is None:
            raise SystemExit('El mapa cita %s!%s («%s») y está VACÍA' % (hoja, coord, etiqueta))
        if etiqueta in mapa:
            raise SystemExit('Etiqueta de mapa repetida: %r' % etiqueta)
        if hasattr(v, 'isoformat'):
            v = v.isoformat()[:10]
        mapa[etiqueta] = {'ref': '%s.xlsx!%s!%s' % (NOMBRE, hoja, coord),
                          'valor': v, 'tipo': tipo}
    if len(mapa) < 25:
        raise SystemExit('El mapa tiene %d etiquetas: el mínimo son 25' % len(mapa))
    with open(os.path.join(C.BUILD_DIR, 'mapa-' + NOMBRE + '.json'), 'w',
              encoding='utf-8') as fh:
        json.dump(mapa, fh, ensure_ascii=False, indent=1)

    return {'ruta': ruta, 'hojas': len(wbf.worksheets), 'formulas': len(motor.REGISTRO),
            'sin_dato': len(sin_dato), 'verdes': n_verdes, 'verdes_vacias': verdes_vacias,
            'notas_legales': len(NOTAS_LEGALES), 'ids_con_nota': len(puestos),
            'mapa': len(mapa),
            'cache': salida_cache.strip().splitlines()[-1] if salida_cache.strip() else ''}


# ==========================================================================
# Filas fijas que se citan entre hojas (se calculan antes de escribir)
# ==========================================================================
# --- Checklist ------------------------------------------------------------
K_CAB = 5
_fila = K_CAB + 1
FASES = []
for _f in sorted(D.FASES_CHECKLIST):
    _items = [i for i, it in enumerate(D.CHECKLIST_LEGAL) if it[0] == _f]
    _sec, _ini = _fila, _fila + 1
    _fin = _ini + len(_items) - 1
    FASES.append((_f, D.FASES_CHECKLIST[_f], _sec, _ini, _fin, _fin + 1, _items))
    _fila = _fin + 3
K_INI, K_FIN = FASES[0][3], FASES[-1][4]
K_EXTRACCION = K_INI            # primera fila: ¿la extracción necesita proyecto?
K_ANEXO = K_INI + 1             # segunda fila: ¿Anexo de la Ley 12/2012?
if not D.CHECKLIST_LEGAL[0][1].startswith('¿La extracción necesita proyecto') \
        or 'Anexo' not in D.CHECKLIST_LEGAL[1][1]:
    raise SystemExit('D13: la 1.ª fila del checklist debe ser la extracción y la 2.ª el Anexo')
K_MAQ = None
for _f, _t, _s, _i, _e, _c, _items in FASES:
    for _pos, _idx in enumerate(_items):
        if 'maquinaria crítica' in D.CHECKLIST_LEGAL[_idx][1]:
            K_MAQ = _i + _pos
if K_MAQ is None:
    raise SystemExit('Falta la fila del checklist que pide la maquinaria crítica')
K_SEC_RES = _fila
(K_TOTAL, K_HECHOS, K_NA, K_AVANCE, K_PREV, K_REAL, K_DESV, K_SINIMP,
 K_CCAA_N) = range(K_SEC_RES + 1, K_SEC_RES + 10)
K_TU_SEC = K_CCAA_N + 2
K_TU_CCAA, K_TU_MUN, K_TU_LECT = K_TU_SEC + 1, K_TU_SEC + 2, K_TU_SEC + 3
K_NOTA = K_TU_LECT + 2
K_LISTAS = K_NOTA + 4

# --- Árbol IAE ------------------------------------------------------------
I = {'sec': 5, 'consumo': 6, 'llevar': 7, 'ferias': 8, 'otros': 9, 'sala': 10,
     'millon': 11, 'sec_f': 13, 'cab': 14, 'ini': 15, 'sec_r': 20, 'n': 21,
     'p6446': 22, 'epigrafes': 23, 'pago': 24, 'pregunta': 25, 'trampa': 26,
     'doble': 27, 'faculta': 28, 'nota': 30, 'listas': 34}
I['fin'] = I['ini'] + len(D.ARBOL_IAE) - 1
SALA_676, SALA_673, SALA_NOSE = 'Chocolatería (676)', 'Café-bar (673)', 'Aún no lo sé'

# --- Régimen y horario ------------------------------------------------------
R = {'sec': 5, 'sup750': 6, 'consumo': 7, 'llevar': 8, 'terraza': 9, 'proyecto': 10,
     'sec_v': 12, 'actividad': 13, 'obras': 14, 'terraza_v': 15,
     'sec_h': 17, 'cab_h': 18, 'cafe': 19, 'choco': 20, 'clas': 22, 'hmin': 23,
     'hmax': 24, 'lectura': 25, 'remite': 26, 'comercio': 27, 'nota': 29, 'listas': 33}
CLAS_CAFE, CLAS_CHOCO = 'Cafetería o bar', 'Chocolatería'

# --- Licencia, humos y gas --------------------------------------------------
G = {'sec': 5, 'formato': 6, 'energia': 7, 'conducto': 8, 'electricos': 9,
     'extintor': 10, 'anios_gas': 11, 'kg_bombona': 12, 'cert_gas': 13,
     'sec_v': 15, 'obra': 16, 'humos': 17, 'gas': 18, 'prox_gas': 19,
     'ext_v': 20, 'riesgo': 21, 'sec_p': 23, 'cab_p': 24, 'p_ini': 25}
PAPELES = [
    ('Licencia de obra del conducto y proyecto técnico', 'obra'),
    ('Limpieza anual del conducto y la campana (ejemplo de Madrid)', 'conducto'),
    ('Certificado de la instalación receptora de gas', 'receptora'),
    ('Inspección periódica de la instalación de gas', 'receptora'),
    ('Extintor de clase F junto a la freidora', 'siempre'),
    ('Autorización de la terraza', 'terraza'),
]
G['p_fin'] = G['p_ini'] + len(PAPELES) - 1
G['n_toca'], G['n_hecho'], G['ver'] = G['p_fin'] + 1, G['p_fin'] + 2, G['p_fin'] + 3
G['nota'], G['listas'] = G['ver'] + 2, G['ver'] + 6
FMT_LOCAL, FMT_DESPACHO, FMT_CASETA = 'Local con sala', 'Despacho para llevar', 'Caseta o feria'
GAS, ELECTRICA = 'Gas', 'Eléctrica'

# --- Registro sanitario -------------------------------------------------------
S = {'sec': 5, 'final': 6, 'otros': 7, 'puesto': 8, 'sec_v': 10, 'regimen': 11,
     'tramite': 12, 'puesto_v': 13, 'autocontrol': 14, 'sec_t': 16, 'cab': 17,
     'caso1': 18, 'nota': 22, 'listas': 25}

# --- Ferias -------------------------------------------------------------------
REQ_FERIA = [
    ('Autorización municipal del puesto: con plazo, sin renovación automática y por '
     'un procedimiento transparente', 'CUN-20'),
    ('Tus datos y la autorización, expuestos de forma visible para el público', 'CUN-20'),
    ('Comprueba qué norma cita tu ordenanza: el real decreto estatal de venta ambulante '
     'de 2010 está derogado desde el 7-8-2021', 'CUN-19'),
    ('Ocupación de dominio público o quiosco: el puesto de churros no siempre va como '
     '«venta ambulante»; hay ordenanzas que lo excluyen de los puestos desmontables',
     'CUN-36'),
    ('Alta en el IAE del puesto: 663.1, o 675 / 674.7 si pones mesas', 'CUN-11 + CUN-13'),
    ('Registro sanitario autonómico: el puesto también es comercio al por menor', 'CUN-21'),
    ('Higiene del puesto: lavado de manos, superficies lisas y lavables, agua potable '
     'caliente y fría, y eliminación higiénica de desechos', 'CUN-22'),
    ('Gas del puesto: la bombona de menos de 15 kg sólo libra de la instalación '
     'receptora a un aparato de utilización móvil', 'CUN-28'),
    ('Ejemplo de Madrid: cocinar en suelo de uso público pide autorización previa y '
     'captación y filtrado a la distancia de la celda B8 de cualquier hueco', 'CUN-25'),
    ('Tasa de la feria: la fija la ordenanza fiscal de cada municipio (el de la nota es '
     'un pueblo pequeño, nunca un orden de magnitud)', 'CUN-37'),
    ('Alérgenos por escrito también en el puesto', 'CHN-34'),
]
FE = {'sec': 5, 'va': 6, 'caduca': 7, 'control': 8, 'dist': 9, 'estado': 10,
      'sec_t': 12, 'cab': 13, 'ini': 14}
FE['fin'] = FE['ini'] + len(REQ_FERIA) - 1
FE['aplican'], FE['pend'], FE['ver'] = FE['fin'] + 1, FE['fin'] + 2, FE['fin'] + 3
FE['nota'], FE['listas'] = FE['ver'] + 2, FE['ver'] + 6

# --- Alérgenos y aceite --------------------------------------------------------
FRITOS = [('Churros y porras', SI), ('Churros rellenos o con cobertura', SI),
          ('Patatas fritas de patata fresca', NO),
          ('Otros fritos (rebozados, empanados, buñuelos de otra masa)', NO)]
ALERGENOS = ('Gluten', 'Leche', 'Huevo', 'Soja', 'Frutos de cáscara')
COLS_AL = ('B', 'C', 'D', 'E', 'F')
#: Lo que se sabe de la carta de El Molinete (SPEC §3.3): gluten siempre en la masa,
#: leche en la taza y en las bebidas con leche. Todo lo demás depende de la etiqueta
#: del preparado que compres, y así se siembra.
ALERGENOS_DEFECTO = {
    'Churros y porras': {'Gluten': SI},
    'Rellenos y especiales': {'Gluten': SI},
    'Chocolate a la taza': {'Leche': SI},
    'Cafés y bebidas': {'Leche': SI},
}
A = {'sec1': 5, 'cab1': 6, 'f_ini': 7}
A['f_fin'] = A['f_ini'] + len(FRITOS) - 1
A['dedicada'] = A['f_fin'] + 1
A['sec2'] = A['dedicada'] + 2
A['cab2'] = A['sec2'] + 1
A['m_ini'] = A['cab2'] + 1
A['m_fin'] = A['m_ini'] + len(D.FAMILIAS) - 1
A['declara'] = A['m_fin'] + 1
A['revisar'] = A['m_fin'] + 2
A['sec3'] = A['revisar'] + 2
A['sg_quiere'], A['sg_umbral'], A['sg_ver'] = A['sec3'] + 1, A['sec3'] + 2, A['sec3'] + 3
A['sec4'] = A['sg_ver'] + 2
(A['ac_patatas'], A['ac_franq'], A['ac_churro'], A['ac_ue'], A['ac_ref'], A['ac_rev'],
 A['ac_a'], A['ac_b']) = range(A['sec4'] + 1, A['sec4'] + 9)
A['sec5'] = A['ac_b'] + 2
A['ac1'] = A['sec5'] + 1
ACEITE_NORMA = [
    ('La norma de calidad de los aceites calentados alcanza a la churrería', 'CUN-01'),
    ('Compuestos polares por debajo del 25 %: el medidor de mano te dice cuándo '
     'cambiar; el 25 % legal se demuestra en laboratorio', 'CUN-02'),
    ('Prohibido añadir sustancias extrañas al aceite y vender el usado para uso '
     'alimentario (el artículo 9 sigue vivo)', 'CUN-03'),
    ('Recogida separada obligatoria del aceite usado desde el 30-06-2022: contrato '
     'con un gestor autorizado y los justificantes guardados', 'CUN-27'),
]
A['nota'] = A['ac1'] + len(ACEITE_NORMA) + 1
A['listas'] = A['nota'] + 4

# --- PRL ----------------------------------------------------------------------
#: Riesgos de la línea de fritura: lista de BUENA PRÁCTICA (supuesto declarado). Las
#: tres filas con norma llevan su id; el resto no cita ninguna (no hay ficha de PRL
#: verificada para este producto).
RIESGOS = [
    ('Quemaduras por salpicadura al cargar la freidora', 'Fritura',
     'Cargar con la cesta o la dosificadora, guantes y delantal resistentes al calor, '
     'nada de agua cerca de la cuba', 2, 3, ''),
    ('Fuego del aceite de la cuba', 'Fritura',
     'Extintor de clase F junto a la freidora y tapa de la cuba a mano; nunca agua',
     1, 3, 'CUN-38'),
    ('Fuego en la campana o el conducto por grasa acumulada', 'Campana y conducto',
     'Filtros a su distancia del foco y limpieza del conducto al menos una vez al año '
     '(ejemplo de Madrid)', 1, 3, 'CUN-26 + CUN-23'),
    ('Fuga de gas o mala combustión', 'Fritura',
     'Instalación receptora certificada e inspeccionada, y ventilación', 1, 3, 'CHN-48'),
    ('Resbalones en suelo con grasa o harina', 'Obrador y fritura',
     'Suelo antideslizante, limpieza durante el turno y calzado cerrado', 2, 2, ''),
    ('Sobreesfuerzos con sacos de harina y garrafas de aceite', 'Almacén',
     'Sacos y garrafas a la altura de la cintura, carro, y los pesados entre dos', 2, 2, ''),
    ('Cortes y atrapamientos en la amasadora y la dosificadora', 'Obrador',
     'Protecciones puestas y máquina parada antes de limpiarla', 1, 3, ''),
    ('Quemaduras al vaciar y filtrar el aceite usado', 'Fritura',
     'Dejar enfriar el aceite antes de vaciarlo y usar un bidón estable', 2, 2, ''),
    ('Calor y humo en la zona de fritura', 'Fritura',
     'Campana encendida todo el turno, pausas e hidratación', 2, 1, ''),
    ('Fatiga por madrugada y turnos partidos', 'Toda la línea',
     'Descanso entre jornadas; las dos figuras de nocturnidad, en el libro de turnos', 2, 2, ''),
]
P = {'alto': 5, 'medio': 6, 'cab': 8, 'ini': 9}
P['fin'] = P['ini'] + len(RIESGOS) - 1
P['n_alta'], P['n_media'], P['pend_alta'], P['ver'] = range(P['fin'] + 2, P['fin'] + 6)
P['nota'], P['listas'] = P['ver'] + 2, P['ver'] + 6

# --- Formación ----------------------------------------------------------------
F_CAB, F_INI = 6, 7
F_FIN = F_INI + len(D.PLANTILLA) + 2      # tres filas libres de más
F_SEC_PAR = F_FIN + 2
F_VALIDEZ, F_CONTROL, F_AVISO = F_SEC_PAR + 1, F_SEC_PAR + 2, F_SEC_PAR + 3
F_SEC_RES = F_SEC_PAR + 5
F_CON, F_CADUCADAS, F_PROXIMAS, F_VER = range(F_SEC_RES + 1, F_SEC_RES + 5)
F_NOTA = F_SEC_RES + 6
F_PIE = F_NOTA + 3
LIBRE_P = '(libre: añade aquí a tu persona)'

# --- Cronograma ---------------------------------------------------------------
GN = {'arranque': 5, 'plazo': 6, 'dias_sem': 7, 'dias_mes': 8, 'sem_mes': 9,
      'centinela': 10, 'objetivo': 11, 'cab': 13, 'ini': 14}
GN['fin'] = GN['ini'] + len(D.GANTT) - 1
GN['sec_mat'] = GN['fin'] + 2
GN['mat_cab'] = GN['sec_mat'] + 1
GN['mat_ini'] = GN['mat_cab'] + 1
GN['mat_fin'] = GN['mat_ini'] + len(D.GANTT) - 1
GN['sec_res'] = GN['mat_fin'] + 2
(GN['duracion'], GN['criticos'], GN['mes_ap'], GN['nombre_mes'], GN['disponibles'],
 GN['margen'], GN['llegas'], GN['maq_critica'], GN['maq_holgura']) = range(
    GN['sec_res'] + 1, GN['sec_res'] + 10)
GN['nota'] = GN['maq_holgura'] + 2
GN['listas'] = GN['nota'] + 4
COLS_MAT = ('C', 'D', 'E', 'F', 'G', 'H', 'I', 'J', 'K', 'L', 'M', 'N', 'O', 'P')
FMT_CENTINELA = '#,##0.0;;;'
if len(COLS_MAT) < len(D.GANTT):
    raise SystemExit('La matriz de dependencias no tiene columnas para %d hitos' % len(D.GANTT))
FILA_H7 = GN['ini'] + [h[0] for h in D.GANTT].index('H7')


# ==========================================================================
# Hoja «Instrucciones»
# ==========================================================================
PASOS = [
    '1. «%s»: los %d trámites en el orden en que hay que hacerlos, en seis fases. '
    'LA PRIMERA FILA es «¿la extracción necesita proyecto?»: si sí, licencia de obra y '
    'técnico aunque la actividad vaya por declaración responsable. La segunda, si tu '
    'actividad está en el Anexo de la Ley 12/2012. Marca el estado, el responsable, lo '
    'que te presupuestan y lo que pagas, y las fechas.' % (H_CHK, len(D.CHECKLIST_LEGAL)),
    '2. «%s»: cinco preguntas y los epígrafes que te tocan (644.6, 676 o 673, 663.1, 675, '
    '419.3), la exención del impuesto y LA PREGUNTA DEL CNAE redactada para tu asesor, '
    'con los dos códigos que hoy se comunican y la trampa del 47.81.' % H_IAE,
    '3. «%s»: por qué régimen va tu actividad, qué pasa con las obras y la terraza, y la '
    'norma del horario del ejemplo de Madrid. Tu hora de apertura NO se pone aquí: se '
    'pone en el libro de temporada y franjas.' % H_REG,
    '4. «%s»: la licencia de obra, la salida de humos del ejemplo de Madrid, el gas '
    '(instalación receptora en el local; la bombona pequeña sólo en la feria) y el '
    'extintor de clase F, con su lista de papeles.' % H_HUM,
    '5. «%s»: registro autonómico o registro general, según a quién vendas.' % H_SAN,
    '6. «%s»: si vas a ferias, los requisitos del puesto. Los euros de cada evento se '
    'calculan en el libro de temporada y ferias.' % H_FER,
    '7. «%s»: qué fríes en la misma freidora, los alérgenos de tu carta por familia, el '
    '«sin gluten», la acrilamida y lo que dice la norma del aceite.' % H_ALE,
    '8. «%s»: los riesgos de la línea de fritura con su prioridad y su medida.' % H_PRL,
    '9. «%s»: una línea por persona con su formación, su fecha y su caducidad.' % H_FOR,
    '10. «%s»: catorce hitos, la ruta crítica y el mes en que abres. AQUÍ, y sólo aquí, '
    'se pone el plazo de entrega de la maquinaria crítica.' % H_GAN,
]

NOTAS_LIBRO = [
    'LA ACTIVIDAD Y LA OBRA SON DOS PAPELES DISTINTOS. Un despacho de churros para llevar '
    '(644.6, hasta 750 metros cuadrados) abre su actividad por declaración responsable, '
    'pero si la salida de humos necesita proyecto -casi siempre que el conducto sube a '
    'cubierta o cruza fachada o patio-, esa obra va con su licencia y su técnico. Con '
    'mesas (676 o 673) la actividad sale del Anexo de la Ley 12/2012 y el modelo lo fija '
    'tu ayuntamiento.',
    'EL CNAE NO TE LO DA ESTA HOJA, Y NO ES UN OLVIDO: la correspondencia oficial entre '
    'el epígrafe del impuesto y el CNAE-2025 no está publicada para la churrería. Lo que '
    'sí te da es la pregunta redactada para tu asesor, con el código de 2025 y el de 2009.',
    'EL HORARIO DE ESTE LIBRO ES UN EJEMPLO DE MADRID y no se extrapola: allí, una '
    'cafetería o bar puede abrir desde las 6:00 y una chocolatería desde las 8:00. Busca '
    'la orden de horarios de tu comunidad.',
    'LA BOMBONA PEQUEÑA NO TE LIBRA DE NADA EN UN LOCAL. La excepción es para un aparato '
    'de utilización móvil, es decir, para el puesto de feria. En un local, gas con '
    'instalación receptora e inspección periódica, o freidora eléctrica.',
    'EL REGISTRO DE FORMACIÓN ES LO QUE SE ACREDITA. La obligación es del empresario: '
    'garantizar la formación de cada persona según su puesto y poder demostrarla.',
    'ESTE LIBRO NO ES UN DICTAMEN JURÍDICO NI SUSTITUYE AL PROYECTO TÉCNICO DE LA '
    'EXTRACCIÓN. Cada dato normativo lleva en su celda un comentario con la norma, el '
    'artículo y el enlace, verificados el %s (o el %s los que vienen de la guía de '
    'chocolatería). Contrástalos con tu ayuntamiento y tu comunidad.'
    % (D.FECHA_VERIFICACION_CUN, D.FECHA_VERIFICACION_CHN),
]

FRONTERA = [
    ('Pack Plantillas APPCC, registro 09 (control del aceite de fritura)',
     'El registro de las mediciones del aceite y de la retirada del aceite usado.',
     'Aquí sólo hay una fila que te recuerda llevarlo. El registro es aquel fichero: se '
     'cita, no se copia.'),
    ('produccion-hora-punta-y-local.xlsx (libro 1 de esta guía)',
     'Los kW de tus freidoras, el riesgo de incendio de la cocina, la extinción '
     'automática y la distancia de los filtros.',
     'Este libro no vuelve a calcular la potencia: te dice qué papeles salen de ella.'),
    ('calculadora-capex-churreria.xlsx (libro 2 de esta guía)',
     'Los euros: proyecto, licencias, obra, extintor y equipamiento, con su IVA.',
     'Los costes de este checklist son los de los TRÁMITES y se piden, no se suponen.'),
    ('temporada-franjas-y-ferias.xlsx (libro 5 de esta guía)',
     'Tu hora de apertura por franja, el semáforo de las 8:00 y los euros de cada feria.',
     'Aquí está la norma del horario y los papeles del puesto; los números, allí.'),
    ('turnos-plantilla-y-madrugada.xlsx (libro 8 de esta guía)',
     'Las dos figuras de nocturnidad, el convenio y el coste de la plantilla.',
     'Aquí sólo están las filas del checklist que te recuerdan hacerlo.'),
    ('Plan de Negocio: Food Truck y Tareas: Food Truck',
     'El caso completo de un puesto móvil, con sus números y sus tareas.',
     'La caseta de feria es aquí un capítulo de papeles, sin caso cifrado.'),
]


def hoja_instrucciones(wb):
    ws = wb.create_sheet(H_INS)
    C.anchos(ws, {'A': 44, 'B': 44, 'C': 50})
    motor.val(ws, 'A1', TITULO)
    ws['A1'].font = Font(bold=True, size=16, color=C.ORO)
    ws.row_dimensions[1].height = 26
    motor.val(ws, 'A2', C.SUBTITULO)
    ws['A2'].font = Font(size=9)
    motor.val(ws, 'A3', 'Para qué sirve: saber qué papel te toca y en qué orden, antes de '
                        'firmar el local y hasta el día que abres.')
    ws['A3'].font = Font(italic=True, size=9)
    fila = 5
    motor.val(ws, 'A%d' % fila, C.NOTA_VERDES, bold=True)
    ws['A%d' % fila].fill = PatternFill('solid', fgColor=motor.VERDE)
    fila += 2
    C.seccion(ws, 'A%d' % fila, 'Instrucciones de uso')
    fila += 1
    for paso in PASOS:
        C.parrafo(ws, fila, paso, 'A', 'C', alto=58)
        fila += 1
    fila += 1

    C.seccion(ws, 'A%d' % fila, 'El orden de relleno de los ocho libros')
    fila += 1
    C.parrafo(ws, fila, 'Rellena los libros en este orden: cada uno usa cifras de los '
              'anteriores, que se copian a mano en su celda verde de cruce. Este libro, el '
              '7, no depende de ninguno: puedes rellenarlo cuando quieras, y conviene '
              'empezarlo antes de firmar el local.', 'A', 'C', alto=40)
    fila += 1
    C.cabecera(ws, fila, [('A', 'Paso'), ('B', 'Libro'), ('C', 'Fichero')], altura=22)
    ws.freeze_panes = None
    fila += 1
    for paso, n in enumerate(D.ORDEN_RELLENO, 1):
        motor.val(ws, 'A%d' % fila, paso, fmt=C.FMT_ENT, align='center')
        motor.val(ws, 'B%d' % fila, 'Libro %d' % n, align='center')
        motor.val(ws, 'C%d' % fila, D.LIBROS[n], bold=(n == LIBRO))
        fila += 1
    fila += 1

    C.seccion(ws, 'A%d' % fila, 'Lo que conviene saber antes de empezar')
    fila += 1
    for texto in NOTAS_LIBRO:
        C.parrafo(ws, fila, texto, 'A', 'C', alto=70)
        fila += 1
    fila += 1

    C.seccion(ws, 'A%d' % fila, 'Qué hace este libro y qué hace otro (para no teclear dos veces)')
    fila += 1
    C.cabecera(ws, fila, [('A', 'Producto o libro'), ('B', 'Qué aporta'),
                          ('C', 'Qué NO hace este libro')], altura=26)
    ws.freeze_panes = None
    fila += 1
    for producto, aporta, no_hace in FRONTERA:
        motor.val(ws, 'A%d' % fila, producto, bold=True, wrap=True)
        motor.val(ws, 'B%d' % fila, aporta, wrap=True)
        motor.val(ws, 'C%d' % fila, no_hace, wrap=True)
        ws.row_dimensions[fila].height = 56
        fila += 1
    fila += 1

    C.seccion(ws, 'A%d' % fila, 'Cada cuánto se usa este libro')
    fila += 1
    C.parrafo(ws, fila,
              'El checklist y el cronograma, UNA VEZ A LA SEMANA desde que decides abrir '
              'hasta que abres. Los árboles (impuesto, régimen y registro), una vez al '
              'principio y otra CADA VEZ QUE CAMBIES DE FORMATO: el día que empieces a '
              'servir churros a la cafetería de al lado o a ir a ferias, cambian tus '
              'papeles. El registro de formación, cada vez que entra alguien. La hoja de '
              'riesgos, cuando cambies de máquina o de distribución.', 'A', 'C', alto=70)
    fila += 2
    C.parrafo(ws, fila, C.NOTA_DESPROTEGER, 'A', 'C', alto=28)
    fila += 2
    fin = C.pie(ws, fila, 'A', 'C')
    C.pagina(ws, apaisado=False, area='A1:C%d' % fin)
    return ws


# ==========================================================================
# Hoja «Checklist Legal (F1-F6)»
# ==========================================================================
def hoja_checklist(wb):
    ws = wb.create_sheet(H_CHK)
    C.anchos(ws, {'A': 13, 'B': 7, 'C': 54, 'D': 20, 'E': 11, 'F': 15, 'G': 14,
                  'H': 13, 'I': 13, 'J': 12, 'K': 13, 'L': 52, 'M': 80})
    C.encabezar(ws, 'Checklist legal y de licencias, en seis fases',
                'Los trámites en el orden en que hay que hacerlos. Las dos primeras filas '
                'son PREGUNTAS: contéstalas antes de firmar el local, porque deciden el '
                'resto. Lo que cambia según tu comunidad va marcado en su columna.',
                col_fin='M')
    C.cabecera(ws, K_CAB,
               [('A', 'Estado'), ('B', 'Fase'), ('C', 'Trámite'), ('D', 'Responsable'),
                ('E', 'Plazo (días)'), ('F', 'Coste previsto (€)'),
                ('G', 'Coste real (€)'), ('H', 'Fecha objetivo'), ('I', 'Fecha hecha'),
                ('J', '¿Cambia por CCAA?'), ('K', 'Tu respuesta'),
                ('L', 'Lo que te toca'), ('M', 'Notas y fuente')])
    refs, fila_listas = C.bloque_listas(
        ws, K_LISTAS,
        [('Estado del trámite', [PENDIENTE, HECHO, NOAPLICA]),
         ('Sí o No', [SI, NO]),
         ('Sí, No o No lo sé', [SI, NO, NOSE]),
         ('Comunidad autónoma', CCAA)], col='A')

    casillas, ccaas, fechas = [], [], []
    for fase, titulo, sec, ini, fin, cont, items in FASES:
        barra_seccion(ws, sec, fase + ' ' + chr(183) + ' ' + titulo, 'BCDEFGHIJKLM')
        for pos, idx in enumerate(items):
            (f, tramite, responsable, plazo, fuente, cambia, nota_txt) = D.CHECKLIST_LEGAL[idx]
            r = ini + pos
            entrada(ws, 'A%d' % r, PENDIENTE, etiqueta=tramite, align='center')
            casillas.append('A%d' % r)
            motor.val(ws, 'B%d' % r, fase, align='center')
            motor.val(ws, 'C%d' % r, tramite, wrap=True, bold=(r in (K_EXTRACCION, K_ANEXO)))
            entrada(ws, 'D%d' % r, responsable, etiqueta='Responsable: ' + tramite, wrap=True)
            if r == K_MAQ:
                # Fuente ÚNICA del plazo de entrega: la celda verde del cronograma.
                motor.f(ws, 'E%d' % r,
                        ie("%s$C$%d*%s$C$%d" % (q(H_GAN), GN['plazo'], q(H_GAN), GN['dias_sem'])),
                        fmt=C.FMT_ENT, align='center')
            else:
                entrada(ws, 'E%d' % r, plazo, fmt=C.FMT_ENT, etiqueta=tramite, align='center')
            entrada(ws, 'F%d' % r, SIN_IMPORTE, etiqueta=tramite)
            entrada(ws, 'G%d' % r, SIN_PAGAR, etiqueta=tramite)
            entrada(ws, 'H%d' % r, FECHA_BASE, fmt=C.FMT_FECHA, etiqueta=tramite)
            entrada(ws, 'I%d' % r, FECHA_BASE, fmt=C.FMT_FECHA, etiqueta=tramite)
            fechas += ['H%d' % r, 'I%d' % r]
            entrada(ws, 'J%d' % r, SI if cambia else NO, etiqueta=tramite, align='center')
            ccaas.append('J%d' % r)
            texto_m = (nota_txt or '') + ('  [Fuente: %s]' % fuente if fuente else '')
            if r == K_MAQ:
                texto_m = ('El plazo de esta fila sale del plazo de entrega de la maquinaria '
                           'crítica, que se pone UNA sola vez, en la hoja «%s» (celda C%d), '
                           'en semanas. Pídelo por escrito al proveedor.  [Fuente: %s]'
                           % (H_GAN, GN['plazo'], fuente))
            C.nota(ws, 'M%d' % r, texto_m.strip())
            ws.row_dimensions[r].height = 44
            notas_de_fuente(ws, ['C%d' % r, 'M%d' % r], fuente)

        motor.val(ws, 'C%d' % cont, 'AVANCE DE LA FASE ' + fase, bold=True)
        motor.f(ws, 'D%d' % cont, ie('COUNTIF(A%d:A%d,"%s")' % (ini, fin, HECHO)),
                fmt=C.FMT_ENT, bold=True)
        motor.f(ws, 'E%d' % cont, ie('COUNTIF(B%d:B%d,"%s")' % (ini, fin, fase)), fmt=C.FMT_ENT)
        motor.f(ws, 'F%d' % cont, ie('SUMIF(F%d:F%d,">=0")' % (ini, fin)), fmt=C.FMT_EUR)
        motor.f(ws, 'G%d' % cont, ie('SUMIF(G%d:G%d,">=0")' % (ini, fin)), fmt=C.FMT_EUR)
        motor.f(ws, 'J%d' % cont,
                ie('IF(E{0}=0,"",(D{0}+COUNTIF(A{1}:A{2},"{3}"))/E{0})'
                   .format(cont, ini, fin, NOAPLICA)), fmt=C.FMT_PCT, bold=True)
        total_fila(ws, cont, 'ABCDEFGHIJKLM')
        motor.semaforo_isnumber(ws, 'J%d' % cont, '$J$%d' % cont, operador='<', umbral='1',
                                bg=motor.CF_AMBAR_BG, fg=motor.CF_AMBAR_FG)
        C.nota(ws, 'M%d' % cont, 'Hechos ' + chr(183) + ' trámites de la fase ' + chr(183)
               + ' coste previsto ' + chr(183) + ' coste real ' + chr(183)
               + ' % de avance (los «N/A» cuentan como resueltos).')

    # --- las dos preguntas de las dos primeras filas (D13, A4) ----------------
    entrada(ws, 'K%d' % K_EXTRACCION, SI, etiqueta='¿La extracción necesita proyecto?',
            align='center')
    C.dv_rango(ws, ['K%d' % K_EXTRACCION], refs['Sí, No o No lo sé'],
               'Contesta Sí, No o No lo sé',
               'Si el conducto sube a cubierta o cruza fachada o patio, casi siempre es Sí. '
               'El Molinete entra en un local sin salida de humos: Sí.')
    motor.f(ws, 'L%d' % K_EXTRACCION,
            ie('IF(K{0}="{s}","Licencia de obra y técnico: el conducto necesita proyecto, '
               'aunque la actividad vaya por declaración responsable",IF(K{0}="{n}","Sin obra '
               'con proyecto: que tu técnico lo deje por escrito antes de firmar","No firmes '
               'el local sin que un técnico mire la salida de humos"))'
               .format(K_EXTRACCION, s=SI, n=NO)), bold=True)
    motor.f(ws, 'K%d' % K_ANEXO,
            ie('IF(OR({iae}$B${c}="{s}",{reg}$B${m}="{s}"),"{n}","{s}")'
               .format(iae=q(H_IAE), c=I['consumo'], reg=q(H_REG), m=R['sup750'], s=SI, n=NO)),
            align='center')
    motor.f(ws, 'L%d' % K_ANEXO,
            ie('IF(K{0}="{s}","Tu actividad está en el Anexo (644.6, hasta 750 m²): va por '
               'declaración responsable; las obras con proyecto, con su licencia aparte",'
               '"Tu actividad sale del Anexo (con mesas o barra, o más de 750 m²): el modelo, '
               'licencia o declaración, lo fija tu ayuntamiento; y las obras con proyecto, con '
               'su licencia aparte")'.format(K_ANEXO, s=SI)), bold=True)
    for r in (K_EXTRACCION, K_ANEXO):
        ws['L%d' % r].alignment = Alignment(vertical='top', wrap_text=True)
        C.destacado(ws, 'L%d' % r)
        ws.row_dimensions[r].height = 66
    motor.regla_expresion(ws, 'L%d' % K_EXTRACCION, '=$K$%d="%s"' % (K_EXTRACCION, SI),
                          bg=motor.CF_AMBAR_BG, fg=motor.CF_AMBAR_FG)

    C.dv_rango(ws, casillas, refs['Estado del trámite'], 'Marca el estado del trámite',
               'Elige %s, %s o %s de la lista del pie de la hoja.' % (PENDIENTE, HECHO, NOAPLICA))
    C.dv_rango(ws, ccaas, refs['Sí o No'], 'Marca Sí o No',
               'Elige «Sí» o «No» de la lista del pie de la hoja.')
    motor.dv_fecha(ws, fechas)
    motor.semaforo_texto(ws, 'A%d:A%d' % (K_INI, K_FIN),
                         ((HECHO, motor.CF_VERDE_BG, motor.CF_VERDE_FG),
                          (NOAPLICA, motor.CF_GRIS_BG, motor.CF_GRIS_FG)))

    # --- resumen -----------------------------------------------------------------
    barra_seccion(ws, K_SEC_RES, 'RESUMEN DE TODO EL EXPEDIENTE', 'BCDEFGHIJKLM')

    def res(fila, etiqueta, formula, fmt=None, nota_txt=None, bold=None):
        motor.val(ws, 'C%d' % fila, etiqueta, bold=bold)
        motor.f(ws, 'D%d' % fila, formula, fmt=fmt, bold=bold)
        if nota_txt:
            C.nota(ws, 'M%d' % fila, nota_txt)

    res(K_TOTAL, 'Trámites del expediente', ie('COUNTIF(B%d:B%d,"F*")' % (K_INI, K_FIN)), C.FMT_ENT)
    res(K_HECHOS, 'Trámites hechos', ie('COUNTIF(A%d:A%d,"%s")' % (K_INI, K_FIN, HECHO)), C.FMT_ENT)
    res(K_NA, 'Trámites que no te aplican (N/A)',
        ie('COUNTIF(A%d:A%d,"%s")' % (K_INI, K_FIN, NOAPLICA)), C.FMT_ENT)
    res(K_AVANCE, 'Avance del expediente (%)',
        ie('IF(D{0}=0,"",(D{1}+D{2})/D{0})'.format(K_TOTAL, K_HECHOS, K_NA)), C.FMT_PCT, bold=True)
    motor.semaforo_isnumber(ws, 'D%d' % K_AVANCE, '$D$%d' % K_AVANCE, operador='<', umbral='1',
                            bg=motor.CF_AMBAR_BG, fg=motor.CF_AMBAR_FG)
    res(K_PREV, 'Coste previsto total del checklist (€)',
        ie('SUMIF(F%d:F%d,">=0")' % (K_INI, K_FIN)), C.FMT_EUR,
        'Sólo los trámites que ya tienen importe. Los euros de la apertura (proyecto, obra, '
        'licencias, equipamiento) con su IVA están en calculadora-capex-churreria.xlsx: aquí '
        'se controla que cada trámite tenga su importe pedido y su fecha.')
    res(K_REAL, 'Coste real total del checklist (€)',
        ie('SUMIF(G%d:G%d,">=0")' % (K_INI, K_FIN)), C.FMT_EUR,
        'Se rellena a medida que pagas. Mientras esté a cero, la desviación no dice nada.')
    res(K_DESV, 'Desviación (real menos previsto, €)', ie('D%d-D%d' % (K_REAL, K_PREV)),
        C.FMT_EUR, bold=True)
    motor.semaforo_isnumber(ws, 'D%d' % K_DESV, '$D$%d' % K_DESV, operador='>', umbral='0')
    res(K_SINIMP, 'Trámites sin importe (a presupuestar)',
        ie('COUNTIF(F%d:F%d,"%s")' % (K_INI, K_FIN, SIN_IMPORTE)), C.FMT_ENT,
        'Dependen de tu ayuntamiento, de tu comunidad, del técnico o del instalador. No se '
        'rellenan a ojo: se piden. Mientras no sea cero, tu presupuesto está incompleto.')
    motor.semaforo_isnumber(ws, 'D%d' % K_SINIMP, '$D$%d' % K_SINIMP, operador='>', umbral='0',
                            bg=motor.CF_AMBAR_BG, fg=motor.CF_AMBAR_FG)
    res(K_CCAA_N, 'Trámites que cambian según tu comunidad autónoma',
        ie('COUNTIF(J%d:J%d,"%s")' % (K_INI, K_FIN, SI)), C.FMT_ENT,
        'Cada uno de éstos hay que confirmarlo con TU comunidad y TU ayuntamiento.')
    total_fila(ws, K_AVANCE, 'CD')
    total_fila(ws, K_DESV, 'CD')

    # --- tu comunidad y tu municipio --------------------------------------------
    barra_seccion(ws, K_TU_SEC, 'TU COMUNIDAD Y TU MUNICIPIO', 'BCDEFGHIJKLM')
    motor.val(ws, 'C%d' % K_TU_CCAA, '¿En qué comunidad autónoma vas a abrir?', bold=True)
    entrada(ws, 'D%d' % K_TU_CCAA, CCAA_EJEMPLO, etiqueta='Comunidad autónoma')
    ws.merge_cells('D%d:F%d' % (K_TU_CCAA, K_TU_CCAA))
    C.dv_rango(ws, ['D%d' % K_TU_CCAA], refs['Comunidad autónoma'], 'Elige tu comunidad',
               'Elige tu comunidad o ciudad autónoma de la lista del pie de la hoja.')
    motor.val(ws, 'C%d' % K_TU_MUN, '¿En qué municipio?', bold=True)
    entrada(ws, 'D%d' % K_TU_MUN, 'Madrid (ejemplo)', etiqueta='Municipio')
    ws.merge_cells('D%d:F%d' % (K_TU_MUN, K_TU_MUN))
    C.nota(ws, 'M%d' % K_TU_MUN, 'Lo pide tu ordenanza de humos, de terrazas y de venta '
                                 'ambulante: cada una es municipal.')
    motor.val(ws, 'C%d' % K_TU_LECT, 'Lo que vale para ti de los ejemplos del libro', bold=True)
    motor.f(ws, 'D%d' % K_TU_LECT,
            ie('IF(D{0}="{m}","Los ejemplos de humos (ordenanza de Madrid) y de horario (orden '
               'de la Comunidad de Madrid) son los de tu comunidad; si tu municipio no es Madrid, '
               'busca además su ordenanza de humos","Los ejemplos de humos y de horario de este '
               'libro son de Madrid: busca la ordenanza de humos de tu municipio y la orden de '
               'horarios de tu comunidad")'.format(K_TU_CCAA, m=CCAA_EJEMPLO)), bold=True)
    ws.merge_cells('D%d:L%d' % (K_TU_LECT, K_TU_LECT))
    ws['D%d' % K_TU_LECT].alignment = Alignment(vertical='top', wrap_text=True)
    ws.row_dimensions[K_TU_LECT].height = 36
    C.destacado(ws, 'D%d' % K_TU_LECT)

    C.parrafo(ws, K_NOTA, 'Los plazos en días son ORIENTATIVOS y casi todos dependen de un '
              'tercero: el ayuntamiento, el técnico, el instalador, la comunidad de '
              'propietarios. El único que controlas del todo es el de pedir las cosas pronto.',
              'A', 'M', alto=34)
    C.parrafo(ws, K_NOTA + 1, 'Nunca «necesitas licencia» o «no la necesitas» sin decir cuál: '
              'la actividad, la obra y la terraza son tres papeles distintos, y las dos '
              'primeras filas de esta hoja te dicen cuáles te tocan.', 'A', 'M', alto=34)
    ws.freeze_panes = 'C%d' % (K_CAB + 1)
    C.pagina(ws, titulos='$%d:$%d' % (K_CAB, K_CAB), area='A1:M%d' % fila_listas)
    return ws


# ==========================================================================
# Hoja «Árbol IAE y CNAE por Formato»
# ==========================================================================
def hoja_iae(wb):
    ws = wb.create_sheet(H_IAE)
    C.anchos(ws, {'A': 44, 'B': 30, 'C': 40, 'D': 40, 'E': 70})
    C.encabezar(ws, 'Árbol IAE y CNAE por formato',
                'Cinco preguntas y los epígrafes que te tocan. El CNAE no sale como dato: '
                'sale como PREGUNTA para tu asesor, con los dos códigos que hoy se '
                'comunican.', col_fin='E')
    refs, fila_listas = C.bloque_listas(
        ws, I['listas'], [('Sí o No', [SI, NO]), ('Tu sala', [SALA_676, SALA_673, SALA_NOSE])],
        col='A')

    barra_seccion(ws, I['sec'], 'LAS PREGUNTAS', 'BCDE')
    preguntas = [
        ('consumo', '¿Tendrás consumo en el local (mesas, barra o terraza)?', SI,
         'El Molinete tiene sala, barra y terraza.'),
        ('llevar', '¿Venderás para llevar (al peso, por docena o por ración)?', SI,
         'El despacho a calle de El Molinete.'),
        ('ferias', '¿Irás a ferias, mercados o casetas?', NO,
         'El Molinete no: en verano elige la carta de verano (decisión 9).'),
        ('otros', '¿Servirás churros a otros negocios (cafeterías, bares, hoteles)?', NO,
         'Si lo haces, cambian tu epígrafe y tu registro sanitario.'),
    ]
    for clave, texto, defecto, nota_txt in preguntas:
        motor.val(ws, 'A%d' % I[clave], texto, wrap=True, bold=True)
        entrada(ws, 'B%d' % I[clave], defecto, etiqueta=texto, align='center')
        C.nota(ws, 'E%d' % I[clave], nota_txt)
    C.dv_rango(ws, ['B%d' % I[k] for k in ('consumo', 'llevar', 'ferias', 'otros', 'millon')],
               refs['Sí o No'], 'Contesta Sí o No', 'Elige «Sí» o «No».')
    motor.val(ws, 'A%d' % I['sala'], '¿Tu sala irá al epígrafe de chocolatería o al de '
                                     'café-bar?', wrap=True, bold=True)
    entrada(ws, 'B%d' % I['sala'], SALA_NOSE, etiqueta='Epígrafe de la sala', align='center')
    C.dv_rango(ws, ['B%d' % I['sala']], refs['Tu sala'], 'Elige el epígrafe de tu sala',
               'Chocolatería (676), café-bar (673) o «Aún no lo sé».')
    C.nota(ws, 'E%d' % I['sala'], 'El ejemplo deja la decisión abierta a propósito: El '
                                  'Molinete se presenta como cafetería para abrir a las 6:00 '
                                  '(hoja de régimen), pero el epígrafe del impuesto se decide '
                                  'con el asesor.')
    motor.val(ws, 'A%d' % I['millon'], '¿Tu cifra de negocio superará 1.000.000 €?', wrap=True,
              bold=True)
    entrada(ws, 'B%d' % I['millon'], NO, etiqueta='Cifra de negocio por encima del millón',
            align='center')
    nota_celda(ws, 'B%d' % I['millon'], 'CHN-73')

    barra_seccion(ws, I['sec_f'], 'LOS FORMATOS Y SUS EPÍGRAFES', 'BCDE')
    C.cabecera(ws, I['cab'], [('A', 'Formato'), ('B', '¿Te toca?'),
                              ('C', 'Epígrafe del impuesto (candidato)'),
                              ('D', 'CNAE para preguntar a tu asesor'), ('E', 'Fuente y nota')])
    nota_celda(ws, 'D%d' % I['cab'], 'CUN-V03')
    claves = ('consumo', 'llevar', 'ferias', 'otros')
    for i, rama in enumerate(D.ARBOL_IAE):
        r = I['ini'] + i
        motor.val(ws, 'A%d' % r, rama['formato'], bold=True, wrap=True)
        motor.f(ws, 'B%d' % r, ie('IF($B${0}="{s}","{s}","{n}")'.format(I[claves[i]], s=SI, n=NO)),
                align='center', bold=True)
        if i == 0:
            motor.f(ws, 'C%d' % r,
                    ie('IF($B${0}="{a}","676 (chocolatería)",IF($B${0}="{b}","673 (cafés y '
                       'bares)","{t}"))'.format(I['sala'], a=SALA_676, b=SALA_673,
                                                t=rama['iae'])))
        else:
            motor.val(ws, 'C%d' % r, rama['iae'], wrap=True)
        motor.val(ws, 'D%d' % r, rama['cnae_a_preguntar'], wrap=True)
        extra = ''
        if 'CUN-16' in rama['fuente']:
            extra = ' La correspondencia oficial IAE-CNAE no está publicada (CUN-16).'
        C.nota(ws, 'E%d' % r, '[Fuente: %s]%s' % (rama['fuente'], extra))
        ws['C%d' % r].alignment = Alignment(vertical='top', wrap_text=True)
        ws.row_dimensions[r].height = 46
        notas_de_fuente(ws, ['A%d' % r, 'C%d' % r], rama['fuente'])
    motor.semaforo_texto(ws, 'B%d:B%d' % (I['ini'], I['fin']),
                         ((SI, motor.CF_VERDE_BG, motor.CF_VERDE_FG),))

    barra_seccion(ws, I['sec_r'], 'LO QUE SALE', 'BCDE')

    def linea(fila, etiqueta, formula, pid=None, nota_txt=None, alto=48):
        motor.val(ws, 'A%d' % fila, etiqueta, wrap=True, bold=True)
        if formula.startswith('='):
            motor.f(ws, 'B%d' % fila, formula, bold=True)
        else:
            motor.val(ws, 'B%d' % fila, formula, wrap=True)
        ws.merge_cells('B%d:D%d' % (fila, fila))
        ws['B%d' % fila].alignment = Alignment(vertical='top', wrap_text=True)
        ws.row_dimensions[fila].height = alto
        if pid:
            nota_celda(ws, 'B%d' % fila, pid)
        if nota_txt:
            C.nota(ws, 'E%d' % fila, nota_txt)

    motor.val(ws, 'A%d' % I['n'], 'Formatos que te tocan', bold=True)
    motor.f(ws, 'B%d' % I['n'], ie('COUNTIF(B%d:B%d,"%s")' % (I['ini'], I['fin'], SI)),
            fmt=C.FMT_ENT, bold=True, align='center')
    linea(I['p6446'], '¿Necesitas el 644.6 además del epígrafe de la sala?',
          ie('IF(AND(B{c}="{s}",B{l}="{s}"),IF(B{sa}="{a}","Con mucha probabilidad sí: el 676 '
             'no lleva la nota que deja vender en el propio establecimiento, así que la venta '
             'para llevar pide además el 644.6 (inferencia: pregúntalo)",IF(B{sa}="{b}",'
             '"Pregúntalo: el 673 deja vender en el propio establecimiento los productos de su '
             'servicio, pero la venta al peso para llevar puede ser otra actividad","Pregúntalo: '
             'depende de si tu sala va al 676 o al 673")),IF(B{l}="{s}","Sin sala, el 644.6 es tu '
             'epígrafe candidato","No vendes para llevar: no aplica"))'
             .format(c=I['consumo'], l=I['llevar'], sa=I['sala'], s=SI, a=SALA_676,
                     b=SALA_673)),
          'CUN-14', 'Una cuota faculta sólo para su actividad: por eso la pregunta.')
    linea(I['epigrafes'], 'Epígrafes candidatos que llevas a tu asesor',
          ie('IF(B{0}="{s}",C{0}&". ","")&IF(B{1}="{s}",C{1}&". ","")&IF(B{2}="{s}",C{2}&". ","")'
             '&IF(B{3}="{s}",C{3}&". ","")'.format(I['ini'], I['ini'] + 1, I['ini'] + 2,
                                                   I['ini'] + 3, s=SI)),
          'CHN-94', 'Una cuota faculta sólo para su actividad: si tienes dos formatos, '
                    'pregunta si necesitas dos altas.')
    linea(I['pago'], '¿Pagas el impuesto?',
          ie('IF(B{0}="{s}","Alta y, con más de un millón de euros de cifra de negocio, la '
             'exención no te alcanza: pregúntale a tu asesor","Alta sí, pago no: están exentos '
             'los dos primeros años de actividad, todas las personas físicas y las sociedades '
             'por debajo de un millón de euros de cifra de negocio")'.format(I['millon'], s=SI)),
          None, 'La nota legal está en la pregunta del millón, más arriba.')
    linea(I['pregunta'], 'LA PREGUNTA DEL CNAE, REDACTADA PARA TU ASESOR',
          ie('"¿Con qué código del CNAE-2025 me doy de alta, y cuál es su equivalente del '
             'CNAE-2009, que también hay que comunicar? Mis candidatos: "&IF(B{0}="{s}",D{0}&"; ","")'
             '&IF(B{1}="{s}",D{1}&"; ","")&IF(B{2}="{s}",D{2}&"; ","")&IF(B{3}="{s}",D{3}&"; ","")'
             '&"y sé que el 47.81 del CNAE-2025 es comercio de vehículos de motor."'
             .format(I['ini'], I['ini'] + 1, I['ini'] + 2, I['ini'] + 3, s=SI)),
          'CUN-15', 'Copia esta celda y envíasela. Las notas explicativas del INE describen '
                    'cada clase por actividad, sin nombrar los churros: por eso se pregunta.',
          alto=74)
    C.destacado(ws, 'B%d' % I['pregunta'])
    linea(I['trampa'], 'La trampa del 47.81',
          'Quien copie el «4781, puestos de venta de alimentos» del CNAE-2009 se encuentra en '
          'el CNAE-2025 con la clase de vehículos de motor. No lo uses.', None,
          'La nota legal está en la pregunta de arriba.')
    linea(I['doble'], 'Dos códigos, no uno',
          'Al darte de alta comunicas el CNAE-2025 y, además, el CNAE-2009 mientras no entre en '
          'vigor la tarifa de primas de accidentes adaptada al de 2025.', 'CHN-74b')
    linea(I['faculta'], 'Una cuota, una actividad',
          'Cada alta faculta sólo para su actividad. Si haces dos cosas (sala y despacho, o '
          'local y feria), pregúntale a tu asesor si son dos altas.', None,
          'La nota legal está en «Epígrafes candidatos».')

    C.parrafo(ws, I['nota'], 'ESTA HOJA NO TE DA TU CNAE NI TU CUOTA, Y NO ES UN OLVIDO. La '
              'correspondencia oficial entre el epígrafe del impuesto y el CNAE-2025 no está '
              'publicada para la churrería, y la cuota no hace falta: casi nadie la paga. Lo '
              'que sí tienes es la pregunta escrita para tu asesor.', 'A', 'E', alto=44)
    C.parrafo(ws, I['nota'] + 1, 'El obrador que sirve a otros negocios (419.3) queda fuera '
              'de este producto: si fabricas para otros dejas de ser minorista. Está aquí sólo '
              'para que sepas dónde está la frontera.', 'A', 'E', alto=34)
    ws.freeze_panes = 'A4'
    C.pagina(ws, apaisado=True, area='A1:E%d' % fila_listas)
    return ws


# ==========================================================================
# Hoja «Régimen de Apertura y Horario»
# ==========================================================================
def hoja_regimen(wb):
    ws = wb.create_sheet(H_REG)
    C.anchos(ws, {'A': 46, 'B': 16, 'C': 16, 'D': 40, 'E': 70})
    C.encabezar(ws, 'Régimen de apertura y horario',
                'Por qué régimen va tu actividad, qué pasa con la obra y la terraza, y la '
                'norma del horario del ejemplo de Madrid. Tu hora de apertura NO se pone '
                'aquí.', col_fin='E')
    refs, fila_listas = C.bloque_listas(
        ws, R['listas'], [('Sí o No', [SI, NO]), ('Clasificación', [CLAS_CAFE, CLAS_CHOCO])],
        col='A')

    barra_seccion(ws, R['sec'], 'LO QUE DECIDE TU RÉGIMEN', 'BCDE')
    motor.val(ws, 'A%d' % R['sup750'], '¿Tu superficie útil de exposición y venta supera '
                                       '750 m²?', wrap=True, bold=True)
    entrada(ws, 'B%d' % R['sup750'], NO, etiqueta='Supera 750 m²', align='center')
    nota_celda(ws, 'B%d' % R['sup750'], 'CUN-35')
    motor.val(ws, 'A%d' % R['consumo'], '¿Consumo en el local? (del árbol de formatos)',
              wrap=True)
    motor.f(ws, 'B%d' % R['consumo'], '=%s$B$%d' % (q(H_IAE), I['consumo']), align='center')
    motor.val(ws, 'A%d' % R['llevar'], '¿Venta para llevar? (del árbol de formatos)', wrap=True)
    motor.f(ws, 'B%d' % R['llevar'], '=%s$B$%d' % (q(H_IAE), I['llevar']), align='center')
    motor.val(ws, 'A%d' % R['terraza'], '¿Terraza o elementos sobre la vía pública?',
              wrap=True, bold=True)
    entrada(ws, 'B%d' % R['terraza'], SI, etiqueta='Terraza', align='center')
    C.nota(ws, 'E%d' % R['terraza'], 'El Molinete tiene una terraza de %.0f m², aparte de los '
                                     '%.0f m² del local.' % (D.dato(D.NEGOCIO, 'm2_terraza'),
                                                             D.dato(D.NEGOCIO, 'm2_total')))
    motor.val(ws, 'A%d' % R['proyecto'], '¿La extracción necesita proyecto? (1.ª fila del '
                                         'checklist)', wrap=True)
    motor.f(ws, 'B%d' % R['proyecto'], '=%s$K$%d' % (q(H_CHK), K_EXTRACCION), align='center')
    C.dv_rango(ws, ['B%d' % R['sup750'], 'B%d' % R['terraza']], refs['Sí o No'],
               'Contesta Sí o No', 'Elige «Sí» o «No».')

    barra_seccion(ws, R['sec_v'], 'TU RÉGIMEN', 'BCDE')

    def linea(fila, etiqueta, formula, pid=None, alto=46):
        motor.val(ws, 'A%d' % fila, etiqueta, wrap=True, bold=True)
        motor.f(ws, 'B%d' % fila, formula, bold=True)
        ws.merge_cells('B%d:E%d' % (fila, fila))
        ws['B%d' % fila].alignment = Alignment(vertical='top', wrap_text=True)
        ws.row_dimensions[fila].height = alto
        if pid:
            nota_celda(ws, 'B%d' % fila, pid)

    linea(R['actividad'], 'Tu actividad',
          ie('IF(B{c}="{s}","Con mesas o barra (676 o 673) tu actividad sale del Anexo de la '
             'Ley 12/2012: el modelo, licencia o declaración responsable, lo fija tu '
             'ayuntamiento",IF(B{m}="{s}","Más de 750 m²: sales del Anexo y manda tu '
             'ayuntamiento",IF(B{l}="{s}","Despacho para llevar (644.6, hasta 750 m²): la '
             'actividad va por declaración responsable","Sin consumo ni venta para llevar: '
             'revisa el árbol de formatos")))'.format(c=R['consumo'], m=R['sup750'],
                                                      l=R['llevar'], s=SI)), 'CHN-49c')
    nota_celda(ws, 'A%d' % R['actividad'], 'CHN-49b')
    linea(R['obras'], 'Tus obras',
          ie('IF(B{0}="{s}","Las obras con proyecto, el conducto a cubierta incluido, van con '
             'su licencia de obra y su técnico, aparte de la actividad",IF(B{0}="{n}","Sin obra '
             'con proyecto según tu respuesta: que tu técnico lo confirme por escrito",'
             '"Pregúntale a un técnico si tu conducto necesita proyecto antes de firmar"))'
             .format(R['proyecto'], s=SI, n=NO)), 'CUN-35')
    linea(R['terraza_v'], 'Tu terraza',
          ie('IF(B{0}="{s}","Autorización de la terraza aparte: la ocupación de la vía pública '
             'no entra en el régimen de la Ley 12/2012 (art. 2.2)","Sin terraza: nada que pedir")'
             .format(R['terraza'], s=SI)))
    C.destacado(ws, 'B%d' % R['actividad'])

    barra_seccion(ws, R['sec_h'], 'EL HORARIO: EL EJEMPLO DE MADRID (NO SE EXTRAPOLA)', 'BCDE')
    C.cabecera(ws, R['cab_h'], [('A', 'Clasificación del establecimiento'), ('B', 'Desde'),
                                ('C', 'Hasta'), ('D', 'Fuente'), ('E', 'Nota')], altura=24)
    ws.freeze_panes = 'A4'
    horas = [(R['cafe'], CLAS_CAFE, 'cafeteria_bar', 2.0),
             (R['choco'], CLAS_CHOCO, 'chocolateria', 1.0)]
    for fila, nombre, clave, hasta in horas:
        desde, fuente, nota_txt = D.HORA_MINIMA_APERTURA_MADRID[clave]
        motor.val(ws, 'A%d' % fila, nombre, bold=True)
        entrada(ws, 'B%d' % fila, desde / 24.0, fmt=FMT_HORA, etiqueta='Desde: ' + nombre,
                align='center')
        entrada(ws, 'C%d' % fila, hasta / 24.0, fmt=FMT_HORA, etiqueta='Hasta: ' + nombre,
                align='center')
        motor.val(ws, 'D%d' % fila, fuente)
        C.nota(ws, 'E%d' % fila, nota_txt)
        nota_celda(ws, 'B%d' % fila, fuente)
    motor.val(ws, 'A%d' % R['clas'], '¿Cómo te clasificas en el ejemplo de Madrid?',
              wrap=True, bold=True)
    entrada(ws, 'B%d' % R['clas'], D.dato(D.NEGOCIO, 'clasificacion_horaria_ejemplo'),
            etiqueta='Clasificación horaria')
    ws.merge_cells('B%d:C%d' % (R['clas'], R['clas']))
    C.dv_rango(ws, ['B%d' % R['clas']], refs['Clasificación'], 'Elige tu clasificación',
               'Cafetería o bar, o chocolatería.')
    motor.val(ws, 'A%d' % R['hmin'], 'Hora mínima de apertura de tu clasificación', bold=True)
    motor.f(ws, 'B%d' % R['hmin'],
            ie('INDEX(B{0}:B{1},MATCH(B{2},A{0}:A{1},0))'.format(R['cafe'], R['choco'], R['clas'])),
            fmt=FMT_HORA, bold=True, align='center')
    motor.val(ws, 'A%d' % R['hmax'], 'Hora máxima de cierre de tu clasificación', bold=True)
    motor.f(ws, 'B%d' % R['hmax'],
            ie('INDEX(C{0}:C{1},MATCH(B{2},A{0}:A{1},0))'.format(R['cafe'], R['choco'], R['clas'])),
            fmt=FMT_HORA, bold=True, align='center')
    linea(R['lectura'], 'Lo que significa para tu desayuno (decisión 8)',
          ie('IF(B{0}="{ch}","Como chocolatería no abres antes de las 8:00 en el ejemplo de '
             'Madrid: si tu desayuno empieza antes, clasifícate como cafetería o bar","Como '
             'cafetería o bar puedes abrir desde las 6:00 en el ejemplo de Madrid")'
             .format(R['clas'], ch=CLAS_CHOCO)), 'CUN-34')
    motor.val(ws, 'A%d' % R['remite'], 'Dónde pones tu hora de apertura', wrap=True, bold=True)
    motor.val(ws, 'B%d' % R['remite'], 'En temporada-franjas-y-ferias.xlsx, hoja «Franjas del '
              'Día»: allí está el semáforo de las 8:00 con tu hora real. Este libro no te pide '
              'el horario.', wrap=True)
    ws.merge_cells('B%d:E%d' % (R['remite'], R['remite']))
    ws.row_dimensions[R['remite']].height = 34
    linea(R['comercio'], 'Si tienes despacho para llevar',
          ie('IF(B{0}="{s}","Tu despacho es comercio al por menor: la Ley de Horarios '
             'Comerciales da libertad horaria a los comercios de menos de 300 m² (inferencia: '
             'confírmalo con tu ayuntamiento; la sala sigue el horario de establecimientos '
             'públicos de tu comunidad)","No vendes para llevar: no aplica")'
             .format(R['llevar'], s=SI)), 'CHN-77')

    C.parrafo(ws, R['nota'], 'EL HORARIO DE MADRID ES UN EJEMPLO Y NO SE EXTRAPOLA a otras '
              'comunidades: cada una tiene su orden de horarios de establecimientos públicos. '
              'Las modificaciones posteriores a la orden de 2022 no se han comprobado.', 'A',
              'E', alto=40)
    C.parrafo(ws, R['nota'] + 1, 'Las horas de la tabla están en verde porque son las de TU '
              'comunidad las que valen: sustitúyelas por las de tu orden de horarios.', 'A',
              'E', alto=30)
    C.pagina(ws, apaisado=True, area='A1:E%d' % fila_listas)
    return ws


# ==========================================================================
# Hoja «Licencia, Humos y Gas»
# ==========================================================================
def hoja_humos(wb):
    ws = wb.create_sheet(H_HUM)
    C.anchos(ws, {'A': 50, 'B': 20, 'C': 18, 'D': 22, 'E': 70})
    C.encabezar(ws, 'Licencia, humos y gas',
                'La licencia de obra del conducto, la salida de humos del ejemplo de Madrid, '
                'el gas y el extintor. Los kW y el riesgo de incendio los calcula el libro 1.',
                col_fin='E')
    refs, fila_listas = C.bloque_listas(
        ws, G['listas'],
        [('Sí o No', [SI, NO]), ('Formato', [FMT_LOCAL, FMT_DESPACHO, FMT_CASETA]),
         ('Energía', [GAS, ELECTRICA]), ('Estado', [PENDIENTE, HECHO, NOAPLICA])], col='A')

    barra_seccion(ws, G['sec'], 'TUS RESPUESTAS', 'BCDE')

    def pregunta(clave, texto, valor, fmt=None, pid=None, nota_txt=None):
        motor.val(ws, 'A%d' % G[clave], texto, wrap=True, bold=True)
        entrada(ws, 'B%d' % G[clave], valor, fmt=fmt, etiqueta=texto, align='center')
        if pid:
            nota_celda(ws, 'B%d' % G[clave], pid)
        if nota_txt:
            C.nota(ws, 'E%d' % G[clave], nota_txt)

    pregunta('formato', '¿Qué formato montas?', FMT_LOCAL,
             nota_txt='El de El Molinete: local con sala, terraza y despacho a calle.')
    pregunta('energia', '¿Tu freidora es de gas o eléctrica?',
             GAS if D.dato(D.PRODUCCION, 'tipo_aparato') == 'gas' else ELECTRICA,
             nota_txt='La freidora de El Molinete es de gas (la de 25 L de la dotación).')
    pregunta('conducto', '¿El conducto subirá a cubierta o cruzará fachada o patio?', SI,
             nota_txt='El Molinete entra en un local sin salida de humos.')
    pregunta('electricos', '¿Todos tus aparatos son eléctricos y cocinan en su interior sin '
                           'transmitir olores?', NO, pid='CUN-24',
             nota_txt='Sólo así, y sin pasar de 4 kW en conjunto, el ejemplo de Madrid no exige '
                      'conducto a cubierta. Una freidora de baño abierto no cocina en su interior.')
    pregunta('extintor', '¿Tienes un extintor de clase F junto a la freidora?', NO,
             nota_txt='Todavía no: es una fila de la fase F4 del checklist.')
    pregunta('anios_gas', 'Años entre inspecciones de la instalación receptora de gas', 5,
             fmt=C.FMT_ENT, pid='CHN-48')
    pregunta('kg_bombona', 'Peso de bombona por debajo del cual no hay instalación receptora '
                           '(kg, sólo con un aparato de utilización móvil)', 15,
             fmt=C.FMT_ENT, pid='CUN-28')
    pregunta('cert_gas', 'Fecha del certificado de tu instalación de gas (o la prevista)',
             FECHA_BASE, fmt=C.FMT_FECHA)
    motor.dv_fecha(ws, ['B%d' % G['cert_gas']])
    C.dv_rango(ws, ['B%d' % G['formato']], refs['Formato'], 'Elige tu formato',
               'Local con sala, despacho para llevar o caseta o feria.')
    C.dv_rango(ws, ['B%d' % G['energia']], refs['Energía'], 'Elige gas o eléctrica',
               'Gas o eléctrica.')
    C.dv_rango(ws, ['B%d' % G[k] for k in ('conducto', 'electricos', 'extintor')],
               refs['Sí o No'], 'Contesta Sí o No', 'Elige «Sí» o «No».')

    barra_seccion(ws, G['sec_v'], 'LO QUE TE TOCA', 'BCDE')

    def linea(clave, etiqueta, formula, fmt=None, pid=None, alto=46):
        fila = G[clave]
        motor.val(ws, 'A%d' % fila, etiqueta, wrap=True, bold=True)
        motor.f(ws, 'B%d' % fila, formula, fmt=fmt, bold=True)
        if not fmt:
            ws.merge_cells('B%d:E%d' % (fila, fila))
            ws['B%d' % fila].alignment = Alignment(vertical='top', wrap_text=True)
        ws.row_dimensions[fila].height = alto
        if pid:
            nota_celda(ws, 'B%d' % fila, pid)

    linea('obra', 'Licencia de obra',
          ie('IF(B{0}="{s}","Licencia de obra y técnico: el conducto necesita proyecto, aunque '
             'la actividad vaya por declaración responsable","Sin conducto que suba o cruce: que '
             'tu técnico confirme por escrito que no hay obra con proyecto")'
             .format(G['conducto'], s=SI)), pid='CUN-35')
    linea('humos', 'Salida de humos (ejemplo de Madrid)',
          ie('IF(B{0}="{s}","Podrías no necesitar conducto a cubierta si además la potencia '
             'conjunta no pasa de 4 kW: compruébalo con tu técnico","Campana con captación, '
             'extracción forzada y filtro de grasas, conducto a cubierta y limpieza al menos una '
             'vez al año")'.format(G['electricos'], s=SI)), pid='CUN-23')
    linea('gas', 'Gas',
          ie('IF(B{e}="{el}","Freidora eléctrica: comprueba la potencia contratada con tu '
             'comercializadora",IF(B{f}="{ca}","Bombona de menos de "&B{kg}&" kg conectada a un '
             'solo aparato de utilización móvil: no es instalación receptora","Instalación '
             'receptora con el certificado del instalador e inspección cada "&B{an}&" años"))'
             .format(e=G['energia'], el=ELECTRICA, f=G['formato'], ca=FMT_CASETA,
                     kg=G['kg_bombona'], an=G['anios_gas'])))
    motor.val(ws, 'A%d' % G['prox_gas'], 'Próxima inspección de tu instalación de gas', bold=True)
    motor.f(ws, 'B%d' % G['prox_gas'],
            ie('IF(OR(B{e}="{el}",B{f}="{ca}"),"",DATE(YEAR(B{c})+B{a},MONTH(B{c}),DAY(B{c})))'
               .format(e=G['energia'], el=ELECTRICA, f=G['formato'], ca=FMT_CASETA,
                       c=G['cert_gas'], a=G['anios_gas'])), fmt=C.FMT_FECHA, bold=True,
            align='center')
    C.nota(ws, 'E%d' % G['prox_gas'], 'Vacía si tu freidora es eléctrica o si es un puesto con '
                                      'bombona pequeña.')
    linea('ext_v', 'Extintor',
          ie('IF(B{0}="{s}","Tienes el agente adecuado al fuego de aceite: la clase F","Falta: el '
             'agente extintor tiene que ser adecuado a la clase de fuego, y el de los aceites y '
             'grasas de cocina es la clase F")'.format(G['extintor'], s=SI)), pid='CUN-38')
    motor.regla_expresion(ws, 'B%d' % G['ext_v'], '=$B$%d="%s"' % (G['extintor'], NO))
    motor.val(ws, 'A%d' % G['riesgo'], 'Potencia y riesgo de incendio', wrap=True, bold=True)
    motor.val(ws, 'B%d' % G['riesgo'], 'Los kW computables, el riesgo de la cocina, la '
              'extinción automática y la distancia de los filtros los calcula '
              'produccion-hora-punta-y-local.xlsx, hoja «Freidoras, Potencia y Riesgo». Aquí no '
              'se vuelven a calcular.', wrap=True)
    ws.merge_cells('B%d:E%d' % (G['riesgo'], G['riesgo']))
    ws.row_dimensions[G['riesgo']].height = 40
    nota_celda(ws, 'B%d' % G['riesgo'], 'CUN-26')

    barra_seccion(ws, G['sec_p'], 'TUS PAPELES DE OBRA, HUMOS Y GAS', 'BCDE')
    C.cabecera(ws, G['cab_p'], [('A', 'Papel'), ('B', '¿Te toca?'), ('C', 'Estado'),
                                ('D', 'Fecha objetivo'), ('E', 'Por qué')], altura=24)
    ws.freeze_panes = 'A4'
    conds = {
        'obra': ('IF(B{0}="{s}","{s}","{n}")'.format(G['conducto'], s=SI, n=NO),
                 'Sale de tu respuesta sobre el conducto.'),
        'conducto': ('IF(B{0}="{s}","{n}","{s}")'.format(G['electricos'], s=SI, n=NO),
                     'Hay conducto salvo que te ampare la excepción de los aparatos eléctricos.'),
        'receptora': ('IF(OR(B{e}="{el}",B{f}="{ca}"),"{n}","{s}")'
                      .format(e=G['energia'], el=ELECTRICA, f=G['formato'], ca=FMT_CASETA,
                              s=SI, n=NO),
                      'Gas en un local o un despacho: instalación receptora.'),
        'siempre': ('IF(B{0}="","","{s}")'.format(G['formato'], s=SI),
                    'Toda freidora: el fuego de aceite es de clase F.'),
        'terraza': ('IF({r}$B${t}="{s}","{s}","{n}")'.format(r=q(H_REG), t=R['terraza'], s=SI,
                                                             n=NO),
                    'Sale de tu respuesta en la hoja de régimen.'),
    }
    estados, fechas = [], []
    for i, (papel, clave) in enumerate(PAPELES):
        r = G['p_ini'] + i
        formula, porque = conds[clave]
        motor.val(ws, 'A%d' % r, papel, wrap=True)
        motor.f(ws, 'B%d' % r, ie(formula), align='center', bold=True)
        entrada(ws, 'C%d' % r, PENDIENTE, etiqueta=papel, align='center')
        estados.append('C%d' % r)
        entrada(ws, 'D%d' % r, FECHA_BASE, fmt=C.FMT_FECHA, etiqueta=papel)
        fechas.append('D%d' % r)
        C.nota(ws, 'E%d' % r, porque)
        ws.row_dimensions[r].height = 30
    C.dv_rango(ws, estados, refs['Estado'], 'Marca el estado', 'Pendiente, Hecho o N/A.')
    motor.dv_fecha(ws, fechas)
    motor.val(ws, 'A%d' % G['n_toca'], 'Papeles que te tocan', bold=True)
    motor.f(ws, 'B%d' % G['n_toca'], ie('COUNTIF(B%d:B%d,"%s")' % (G['p_ini'], G['p_fin'], SI)),
            fmt=C.FMT_ENT, bold=True, align='center')
    motor.val(ws, 'A%d' % G['n_hecho'], 'De ellos, hechos', bold=True)
    motor.f(ws, 'B%d' % G['n_hecho'],
            '=SUMPRODUCT(--(B{a}:B{b}="{s}"),--(C{a}:C{b}="{h}"))'
            .format(a=G['p_ini'], b=G['p_fin'], s=SI, h=HECHO), fmt=C.FMT_ENT, bold=True,
            align='center')
    motor.val(ws, 'A%d' % G['ver'], 'VEREDICTO', bold=True)
    motor.f(ws, 'B%d' % G['ver'],
            ie('IF(B{t}=0,"No te toca ningún papel de esta hoja",IF(B{h}<B{t},"Te faltan papeles '
               'de obra, humos o gas: "&(B{t}-B{h})&" pendientes","Papeles de obra, humos y gas en '
               'regla"))'.format(t=G['n_toca'], h=G['n_hecho'])), bold=True)
    ws.merge_cells('B%d:E%d' % (G['ver'], G['ver']))
    C.destacado(ws, 'B%d' % G['ver'])
    total_fila(ws, G['n_toca'], 'AB')

    C.parrafo(ws, G['nota'], 'Lo que fija el proyecto técnico -caudal, sección y trazado del '
              'conducto- no lo fija este libro. Los requisitos de humos son los de la ordenanza '
              'de Madrid, el único municipio verificado para este producto: busca la tuya.', 'A',
              'E', alto=40)
    C.parrafo(ws, G['nota'] + 1, 'La bombona pequeña NO sirve para la freidora fija de un local '
              'ni de un despacho: la excepción es para un aparato de utilización móvil, es '
              'decir, para el puesto de feria.', 'A', 'E', alto=34)
    C.pagina(ws, apaisado=True, area='A1:E%d' % fila_listas)
    return ws


# ==========================================================================
# Hoja «Árbol de Registro Sanitario»
# ==========================================================================
CASOS_REGISTRO = (
    ('Vendes al consumidor final en tu local, tu despacho o tu web',
     'Comunicación o declaración responsable al registro de tu comunidad autónoma',
     'El comercio al por menor está excluido del registro general (RGSEAA). Tener obrador '
     'de masa a la vista no te saca del régimen minorista.'),
    ('Tienes además un puesto de feria o una caseta',
     'El mismo registro autonómico para el puesto, y los requisitos de higiene del puesto',
     'El local ambulante o provisional también es comercio al por menor.'),
    ('Sirves churros a una cafetería o un bar de otro titular',
     'Pregunta a tu comunidad: puede tocarte otra inscripción',
     'La guía de la Comunidad de Madrid pone de ejemplo de actividad restringida «la '
     'churrería o pastelería que sirve a la cafetería».'),
)


def hoja_sanitario(wb):
    ws = wb.create_sheet(H_SAN)
    C.anchos(ws, {'A': 56, 'B': 34, 'C': 90})
    C.encabezar(ws, '¿Registro autonómico o registro general?',
                'Tres preguntas y un veredicto. Dos de ellas salen del árbol de formatos: '
                'si cambias de formato allí, cambia aquí.', col_fin='C')
    refs, fila_listas = C.bloque_listas(ws, S['listas'], [('Sí o No', [SI, NO])], col='A')
    barra_seccion(ws, S['sec'], 'LAS PREGUNTAS', 'BC')
    motor.val(ws, 'A%d' % S['final'], '¿Vendes al consumidor final en tu propio '
                                      'establecimiento?', wrap=True, bold=True)
    entrada(ws, 'B%d' % S['final'], SI, etiqueta='Consumidor final', align='center')
    C.dv_rango(ws, ['B%d' % S['final']], refs['Sí o No'], 'Contesta Sí o No', 'Elige «Sí» o «No».')
    motor.val(ws, 'A%d' % S['otros'], '¿Sirves churros a otros negocios? (del árbol de formatos)',
              wrap=True)
    motor.f(ws, 'B%d' % S['otros'], '=%s$B$%d' % (q(H_IAE), I['otros']), align='center')
    motor.val(ws, 'A%d' % S['puesto'], '¿Vendes en un puesto, caseta o vehículo? (del árbol de '
                                       'formatos)', wrap=True)
    motor.f(ws, 'B%d' % S['puesto'], '=%s$B$%d' % (q(H_IAE), I['ferias']), align='center')

    barra_seccion(ws, S['sec_v'], 'EL VEREDICTO', 'BC')

    def linea(fila, etiqueta, formula, pid=None):
        motor.val(ws, 'A%d' % fila, etiqueta, wrap=True, bold=True)
        if formula.startswith('='):
            motor.f(ws, 'B%d' % fila, formula, bold=True)
        else:
            motor.val(ws, 'B%d' % fila, formula, wrap=True)
        ws.merge_cells('B%d:C%d' % (fila, fila))
        ws['B%d' % fila].alignment = Alignment(vertical='top', wrap_text=True)
        ws.row_dimensions[fila].height = 46
        if pid:
            nota_celda(ws, 'B%d' % fila, pid)

    linea(S['regimen'], 'Tu régimen sanitario',
          ie('IF(B{o}="{s}","Sirves a otro negocio: la guía de la Comunidad de Madrid lo pone de '
             'ejemplo de actividad restringida. Pregunta a tu comunidad qué inscripción te toca",'
             'IF(B{f}="{s}","Comercio al por menor: excluido del registro general (RGSEAA); te '
             'inscribes en el de tu comunidad autónoma","Revísalo: si no vendes al consumidor '
             'final, no eres un minorista"))'.format(o=S['otros'], f=S['final'], s=SI)), 'CUN-30')
    linea(S['tramite'], 'El trámite que te toca',
          ie('IF(B{f}="{s}","Comunicación o declaración responsable al registro de tu comunidad '
             'autónoma: no habilita por sí sola, y cómo funciona cambia en cada comunidad","")'
             .format(f=S['final'], s=SI)), 'CHN-39')
    linea(S['puesto_v'], 'Si vas a ferias',
          ie('IF(B{0}="{s}","El puesto, la caseta o el vehículo también es comercio al por '
             'menor: mismo registro autonómico y los requisitos de higiene del puesto (hoja de '
             'ferias)","No vas a ferias: nada más")'.format(S['puesto'], s=SI)), 'CUN-21')
    linea(S['autocontrol'], 'El autocontrol',
          'Plan de autocontrol basado en el APPCC, que puede ser simplificado, con una persona '
          'responsable con nombre. Los registros, en el Pack APPCC (el del aceite es el 09).',
          'CHN-32')
    C.destacado(ws, 'B%d' % S['regimen'])

    barra_seccion(ws, S['sec_t'], 'LOS TRES CASOS, PARA QUE VEAS DÓNDE CAES', 'BC')
    C.cabecera(ws, S['cab'], [('A', 'Tu situación'), ('B', 'Lo que te toca'), ('C', 'Por qué')],
               altura=24)
    ws.freeze_panes = 'A4'
    for i, (sit, toca, porque) in enumerate(CASOS_REGISTRO):
        r = S['caso1'] + i
        motor.val(ws, 'A%d' % r, sit, wrap=True)
        motor.val(ws, 'B%d' % r, toca, wrap=True)
        motor.val(ws, 'C%d' % r, porque, wrap=True)
        ws.row_dimensions[r].height = 48
    C.parrafo(ws, S['nota'], 'Si vendes al consumidor final desde tu propio establecimiento, tu '
              'sitio es el registro de tu comunidad autónoma, no el estatal. Lo que cambia la '
              'cosa es servir a OTRO negocio: el obrador que sirve a otros queda fuera de este '
              'producto.', 'A', 'C', alto=40)
    C.pagina(ws, apaisado=False, area='A1:C%d' % fila_listas)
    return ws


# ==========================================================================
# Hoja «Ferias y Venta Ambulante»
# ==========================================================================
def hoja_ferias(wb):
    ws = wb.create_sheet(H_FER)
    C.anchos(ws, {'A': 66, 'B': 14, 'C': 14, 'D': 20, 'E': 70})
    C.encabezar(ws, 'Ferias y venta ambulante',
                'Los papeles del puesto, la caseta o el remolque. Los euros de cada evento '
                '(tasa, afluencia y punto muerto) están en temporada-franjas-y-ferias.xlsx.',
                col_fin='E')
    refs, fila_listas = C.bloque_listas(ws, FE['listas'],
                                        [('Estado', [PENDIENTE, HECHO, NOAPLICA])], col='A')
    barra_seccion(ws, FE['sec'], 'TU PUESTO', 'BCDE')
    motor.val(ws, 'A%d' % FE['va'], '¿Vas a ferias? (del árbol de formatos)', wrap=True, bold=True)
    motor.f(ws, 'B%d' % FE['va'], '=%s$B$%d' % (q(H_IAE), I['ferias']), align='center', bold=True)
    motor.val(ws, 'A%d' % FE['caduca'], 'Fecha en que caduca tu autorización municipal', bold=True)
    entrada(ws, 'B%d' % FE['caduca'], datetime.date(2027, 3, 31), fmt=C.FMT_FECHA,
            etiqueta='Caducidad de la autorización')
    nota_celda(ws, 'B%d' % FE['caduca'], 'CUN-20')
    motor.val(ws, 'A%d' % FE['control'], 'Fecha de control (ponla en «hoy»)', bold=True)
    entrada(ws, 'B%d' % FE['control'], FECHA_BASE, fmt=C.FMT_FECHA, etiqueta='Fecha de control')
    motor.dv_fecha(ws, ['B%d' % FE['caduca'], 'B%d' % FE['control']])
    motor.val(ws, 'A%d' % FE['dist'], 'Distancia mínima a cualquier hueco si cocinas con aceite '
                                      'en suelo de uso público (m, ejemplo de Madrid)', wrap=True,
              bold=True)
    entrada(ws, 'B%d' % FE['dist'], 15, fmt=C.FMT_ENT, etiqueta='Distancia a hueco receptor',
            align='center')
    nota_celda(ws, 'B%d' % FE['dist'], 'CUN-25')
    motor.val(ws, 'A%d' % FE['estado'], 'Estado de tu autorización', bold=True)
    motor.f(ws, 'B%d' % FE['estado'],
            ie('IF(B{v}<>"{s}","No vas a ferias",IF(B{c}<B{k},"CADUCADA: la autorización no es '
               'indefinida ni se renueva sola","Vigente"))'.format(v=FE['va'], c=FE['caduca'],
                                                                 k=FE['control'], s=SI)),
            bold=True)
    ws.merge_cells('B%d:E%d' % (FE['estado'], FE['estado']))
    motor.regla_expresion(ws, 'B%d' % FE['estado'],
                          '=AND($B$%d="%s",ISNUMBER($B$%d),ISNUMBER($B$%d),$B$%d<$B$%d)'
                          % (FE['va'], SI, FE['caduca'], FE['control'], FE['caduca'], FE['control']))

    barra_seccion(ws, FE['sec_t'], 'LOS REQUISITOS DEL PUESTO', 'BCDE')
    C.cabecera(ws, FE['cab'], [('A', 'Requisito'), ('B', '¿Te aplica?'), ('C', 'Estado'),
                               ('D', 'Fuente'), ('E', 'Nota')], altura=24)
    ws.freeze_panes = 'A4'
    estados = []
    for i, (req, fuente) in enumerate(REQ_FERIA):
        r = FE['ini'] + i
        motor.val(ws, 'A%d' % r, req, wrap=True)
        motor.f(ws, 'B%d' % r, ie('IF($B${0}="{s}","{s}","{n}")'.format(FE['va'], s=SI, n=NO)),
                align='center', bold=True)
        entrada(ws, 'C%d' % r, PENDIENTE, etiqueta=req[:60], align='center')
        estados.append('C%d' % r)
        motor.val(ws, 'D%d' % r, fuente)
        ws.row_dimensions[r].height = 44
        ids = notas_de_fuente(ws, ['A%d' % r, 'D%d' % r], fuente)
        if fuente == 'CUN-19':
            C.nota(ws, 'E%d' % r, 'Lo sustituyó otro real decreto en 2021. Si tu ordenanza cita '
                                  'sus artículos, pregúntale al ayuntamiento qué norma aplica hoy.')
        elif fuente == 'CUN-37':
            C.nota(ws, 'E%d' % r, 'Los euros de la tasa de cada evento se ponen en el libro de '
                                  'temporada y ferias, no aquí.')
    C.dv_rango(ws, estados, refs['Estado'], 'Marca el estado', 'Pendiente, Hecho o N/A.')
    motor.val(ws, 'A%d' % FE['aplican'], 'Requisitos que te aplican', bold=True)
    motor.f(ws, 'B%d' % FE['aplican'], ie('COUNTIF(B%d:B%d,"%s")' % (FE['ini'], FE['fin'], SI)),
            fmt=C.FMT_ENT, bold=True, align='center')
    motor.val(ws, 'A%d' % FE['pend'], 'De ellos, pendientes', bold=True)
    motor.f(ws, 'B%d' % FE['pend'],
            '=SUMPRODUCT(--(B{a}:B{b}="{s}"),--(C{a}:C{b}="{p}"))'
            .format(a=FE['ini'], b=FE['fin'], s=SI, p=PENDIENTE), fmt=C.FMT_ENT, bold=True,
            align='center')
    motor.semaforo_isnumber(ws, 'B%d' % FE['pend'], '$B$%d' % FE['pend'], '>', '0',
                            bg=motor.CF_AMBAR_BG, fg=motor.CF_AMBAR_FG)
    motor.val(ws, 'A%d' % FE['ver'], 'VEREDICTO', bold=True)
    motor.f(ws, 'B%d' % FE['ver'],
            ie('IF(B{v}<>"{s}","No vas a ferias: esta hoja no te aplica. Si algún día vas, '
               'cambia la respuesta en el árbol de formatos",IF(B{p}>0,"Te faltan papeles del '
               'puesto: "&B{p}&" pendientes","Papeles del puesto en regla"))'
               .format(v=FE['va'], p=FE['pend'], s=SI)), bold=True)
    ws.merge_cells('B%d:E%d' % (FE['ver'], FE['ver']))
    C.destacado(ws, 'B%d' % FE['ver'])
    C.parrafo(ws, FE['nota'], 'La caseta es aquí un capítulo de papeles, sin caso cifrado. El '
              'caso completo de un puesto móvil, con sus números y sus tareas, está en el Plan '
              'de Negocio Food Truck y en el Kit de Tareas Food Truck.', 'A', 'E', alto=34)
    C.parrafo(ws, FE['nota'] + 1, 'Cada ayuntamiento tiene su ordenanza de venta ambulante, de '
              'ocupación de la vía pública y su ordenanza fiscal: los requisitos de esta hoja son '
              'la lista de lo que hay que preguntar, no la respuesta de tu municipio.', 'A', 'E',
              alto=34)
    C.pagina(ws, apaisado=True, area='A1:E%d' % fila_listas)
    return ws


# ==========================================================================
# Hoja «Alérgenos y Aceite Compartido»
# ==========================================================================
def hoja_alergenos(wb):
    ws = wb.create_sheet(H_ALE)
    C.anchos(ws, {'A': 50, 'B': 16, 'C': 16, 'D': 16, 'E': 16, 'F': 18, 'G': 70})
    C.encabezar(ws, 'Alérgenos y aceite compartido',
                'Lo que fríes en la misma freidora, los alérgenos de tu carta por familia, el '
                '«sin gluten», la acrilamida y lo que dice la norma del aceite.', col_fin='G')
    refs, fila_listas = C.bloque_listas(
        ws, A['listas'], [('Sí o No', [SI, NO]), ('Alérgeno', [SI, NO, REVISA])], col='A')

    barra_seccion(ws, A['sec1'], '1. QUÉ FRÍES EN LA FREIDORA DE CHURROS', 'BCDEFG')
    C.cabecera(ws, A['cab1'], [('A', 'Lo que fríes'), ('B', '¿En la misma freidora?'),
                               ('G', 'Nota')], altura=24)
    ws.freeze_panes = 'A4'
    v_frito = []
    for i, (que, defecto) in enumerate(FRITOS):
        r = A['f_ini'] + i
        motor.val(ws, 'A%d' % r, que, wrap=True)
        entrada(ws, 'B%d' % r, defecto, etiqueta='Fríes ' + que, align='center')
        v_frito.append('B%d' % r)
    C.dv_rango(ws, v_frito, refs['Sí o No'], 'Contesta Sí o No', 'Elige «Sí» o «No».')
    C.nota(ws, 'G%d' % A['f_ini'], 'El Molinete fríe sólo churros y porras en su freidora; los '
                                   'rellenos y la cobertura van después de freír.')
    fila_patatas = A['f_ini'] + 2
    motor.val(ws, 'A%d' % A['dedicada'], '¿Tu freidora es dedicada o compartida? (decisión 11)',
              bold=True, wrap=True)
    motor.f(ws, 'B%d' % A['dedicada'],
            ie('IF(COUNTIF(B{0}:B{1},"{s}")>0,"Compartida: el aceite arrastra los alérgenos de '
               'todo lo que fríes en él. Freidora aparte, o limpieza antes de cambiar de producto",'
               '"Dedicada al churro y a la porra: el aceite sólo lleva lo que lleva tu masa")'
               .format(fila_patatas, A['f_fin'], s=SI)), bold=True)
    ws.merge_cells('B%d:G%d' % (A['dedicada'], A['dedicada']))
    ws['B%d' % A['dedicada']].alignment = Alignment(vertical='top', wrap_text=True)
    ws.row_dimensions[A['dedicada']].height = 36
    nota_celda(ws, 'B%d' % A['dedicada'], 'CHN-33')
    C.destacado(ws, 'B%d' % A['dedicada'])

    barra_seccion(ws, A['sec2'], '2. LOS ALÉRGENOS DE TU CARTA, POR FAMILIA', 'BCDEFG')
    C.cabecera(ws, A['cab2'], [('A', 'Familia de la carta')]
               + [(COLS_AL[k], ALERGENOS[k]) for k in range(len(ALERGENOS))]
               + [('G', 'Nota')], altura=24)
    nota_celda(ws, 'A%d' % A['cab2'], 'CHN-34')
    celdas = []
    for i, fam in enumerate(D.FAMILIAS):
        r = A['m_ini'] + i
        motor.val(ws, 'A%d' % r, fam, bold=True)
        for k, al in enumerate(ALERGENOS):
            v = ALERGENOS_DEFECTO.get(fam, {}).get(al, REVISA)
            entrada(ws, '%s%d' % (COLS_AL[k], r), v, etiqueta='%s: %s' % (fam, al),
                    align='center')
            celdas.append('%s%d' % (COLS_AL[k], r))
    C.dv_rango(ws, celdas, refs['Alérgeno'], 'Sí, No o Revisa la etiqueta',
               'Si no lo sabes, deja «Revisa la etiqueta» hasta leer la del preparado.')
    motor.semaforo_texto(ws, '%s%d:%s%d' % (COLS_AL[0], A['m_ini'], COLS_AL[-1], A['m_fin']),
                         ((SI, motor.CF_ROJO_BG, motor.CF_ROJO_FG),
                          (REVISA, motor.CF_AMBAR_BG, motor.CF_AMBAR_FG)))
    C.nota(ws, 'G%d' % A['m_ini'], 'Gluten siempre en la masa de trigo. Lo demás depende de la '
                                   'etiqueta del mix y del preparado que compres.')
    C.nota(ws, 'G%d' % (A['m_ini'] + 2), 'Leche en la taza: compruébalo en la etiqueta de tu '
                                         'preparado.')
    motor.val(ws, 'A%d' % A['declara'], '¿Lo declaras por escrito en la carta o el cartel?',
              bold=True, wrap=True)
    for k in range(len(ALERGENOS)):
        col = COLS_AL[k]
        motor.f(ws, '%s%d' % (col, A['declara']),
                ie('IF(COUNTIF({c}{a}:{c}{b},"{s}")>0,"{s}","{n}")'
                   .format(c=col, a=A['m_ini'], b=A['m_fin'], s=SI, n=NO)),
                align='center', bold=True)
    motor.val(ws, 'A%d' % A['revisar'], 'Casillas pendientes de leer en la etiqueta', bold=True)
    motor.f(ws, 'B%d' % A['revisar'],
            ie('COUNTIF(%s%d:%s%d,"%s")' % (COLS_AL[0], A['m_ini'], COLS_AL[-1], A['m_fin'], REVISA)),
            fmt=C.FMT_ENT, bold=True, align='center')
    motor.semaforo_isnumber(ws, 'B%d' % A['revisar'], '$B$%d' % A['revisar'], '>', '0',
                            bg=motor.CF_AMBAR_BG, fg=motor.CF_AMBAR_FG)
    C.nota(ws, 'G%d' % A['revisar'], 'Mientras no sea cero, tu carta de alérgenos está '
                                     'incompleta. La matriz por elaboración está en el Pack APPCC.')
    total_fila(ws, A['declara'], 'ABCDEF')

    barra_seccion(ws, A['sec3'], '3. EL «SIN GLUTEN»', 'BCDEFG')
    motor.val(ws, 'A%d' % A['sg_quiere'], '¿Quieres anunciar algo «sin gluten»?', bold=True)
    entrada(ws, 'B%d' % A['sg_quiere'], NO, etiqueta='Anunciar sin gluten', align='center')
    C.dv_rango(ws, ['B%d' % A['sg_quiere']], refs['Sí o No'], 'Contesta Sí o No',
               'Elige «Sí» o «No».')
    motor.val(ws, 'A%d' % A['sg_umbral'], 'Gluten máximo para decir «sin gluten» (mg/kg)',
              bold=True)
    entrada(ws, 'B%d' % A['sg_umbral'], 20, fmt=C.FMT_ENT, etiqueta='Umbral sin gluten',
            align='center')
    nota_celda(ws, 'B%d' % A['sg_umbral'], 'CHN-37')
    motor.val(ws, 'A%d' % A['sg_ver'], 'Lo que significa', bold=True)
    motor.f(ws, 'B%d' % A['sg_ver'],
            ie('IF(B{q}="{s}","Casi nunca puedes: con una freidora que fríe churros de trigo el '
               'aceite contamina. Necesitarías freidora y elaboración aparte, y un análisis por '
               'debajo de "&B{u}&" mg/kg en el producto tal como lo vendes","No lo anuncias: '
               'correcto. «Sin gluten» no es un argumento de carta en una churrería")'
               .format(q=A['sg_quiere'], u=A['sg_umbral'], s=SI)), bold=True)
    ws.merge_cells('B%d:G%d' % (A['sg_ver'], A['sg_ver']))
    ws['B%d' % A['sg_ver']].alignment = Alignment(vertical='top', wrap_text=True)
    ws.row_dimensions[A['sg_ver']].height = 40

    barra_seccion(ws, A['sec4'], '4. LA ACRILAMIDA', 'BCDEFG')

    def linea(fila, etiqueta, contenido, pid=None, verde=None, alto=40):
        motor.val(ws, 'A%d' % fila, etiqueta, wrap=True, bold=True)
        if verde is not None:
            entrada(ws, 'B%d' % fila, verde, etiqueta=etiqueta, align='center')
        elif contenido.startswith('='):
            motor.f(ws, 'B%d' % fila, contenido, bold=True)
            ws.merge_cells('B%d:G%d' % (fila, fila))
        else:
            motor.val(ws, 'B%d' % fila, contenido, wrap=True)
            ws.merge_cells('B%d:G%d' % (fila, fila))
        ws['B%d' % fila].alignment = Alignment(vertical='top', wrap_text=True)
        ws.row_dimensions[fila].height = alto
        if pid:
            nota_celda(ws, 'B%d' % fila, pid)

    linea(A['ac_patatas'], '¿Fríes patatas fritas de patata fresca? (de la tabla 1)',
          '=B%d' % fila_patatas)
    linea(A['ac_franq'], '¿Operas bajo marca o licencia siguiendo las instrucciones de quien te '
                         'suministra de forma centralizada?', None, verde=NO, alto=46)
    C.dv_rango(ws, ['B%d' % A['ac_franq']], refs['Sí o No'], 'Contesta Sí o No',
               'Elige «Sí» o «No».')
    linea(A['ac_churro'], 'El churro',
          'Buena práctica, no obligación expresa: el churro no está nombrado en la lista del '
          'reglamento de la acrilamida, que no define «bollería». No se puede afirmar ni que '
          'entre ni que quede fuera.', 'CUN-04')
    linea(A['ac_ue'], 'La recomendación de la UE',
          'Una recomendación de la Comisión, no vinculante, incluye los churros en su lista de '
          'alimentos en los que se debe vigilar la acrilamida.', 'CUN-06')
    linea(A['ac_ref'], 'Valor de referencia',
          'Un informe oficial recoge que no hay valor de referencia para las masas fritas como '
          'los churros: la autoridad ya mira el churro, pero no fija una obligación.', 'CUN-07')
    linea(A['ac_rev'], 'Lo que se mueve',
          'La Comisión revisa el marco de la acrilamida (estado de consultas). Vuelve a mirarlo '
          'en el anexo normativo de la guía antes de cada temporada.', 'CUN-08')
    linea(A['ac_a'], 'Parte A del anexo II',
          ie('IF(B{0}="{s}","Te toca la parte A: las medidas de mitigación de las patatas fritas, '
             'porque las fríes","No fríes patatas: la parte A no te alcanza por ellas")'
             .format(A['ac_patatas'], s=SI)), 'CUN-05')
    linea(A['ac_b'], 'Parte B del anexo II',
          ie('IF(AND(B{f}="{s}",B{p}="{s}"),"Parte B, además: operas siguiendo las '
             'instrucciones de quien te suministra de forma centralizada (art. 2.3)","La parte B '
             'no te alcanza: sólo aplica a quien opera bajo marca o licencia siguiendo las '
             'instrucciones de un suministro centralizado (art. 2.3)")'
             .format(f=A['ac_franq'], p=A['ac_patatas'], s=SI)))

    barra_seccion(ws, A['sec5'], '5. EL ACEITE: LO QUE DICE LA NORMA', 'BCDEFG')
    for i, (texto, pid) in enumerate(ACEITE_NORMA):
        r = A['ac1'] + i
        motor.val(ws, 'A%d' % r, texto, wrap=True)
        ws.merge_cells('A%d:F%d' % (r, r))
        ws.row_dimensions[r].height = 32
        nota_celda(ws, 'A%d' % r, pid)
        motor.val(ws, 'G%d' % r, '[Fuente: %s]' % pid)
    C.parrafo(ws, A['nota'], 'El registro de las mediciones del aceite y de la retirada del '
              'usado se lleva en el Pack APPCC, registro 09: aquí se cita, no se copia. El coste '
              'del aceite y cuándo cambiarlo, en aceite-de-fritura-coste-y-cambio.xlsx.', 'A',
              'G', alto=34)
    C.parrafo(ws, A['nota'] + 1, 'Cada casilla de la tabla de alérgenos se rellena con la '
              'etiqueta del producto que compras en la mano: el mismo mix de churros cambia de '
              'composición entre marcas.', 'A', 'G', alto=30)
    C.pagina(ws, apaisado=True, area='A1:G%d' % fila_listas)
    return ws


# ==========================================================================
# Hoja «PRL de la Fritura»
# ==========================================================================
def hoja_prl(wb):
    ws = wb.create_sheet(H_PRL)
    C.anchos(ws, {'A': 5, 'B': 40, 'C': 16, 'D': 50, 'E': 12, 'F': 12, 'G': 10, 'H': 12,
                  'I': 12, 'J': 18, 'K': 50})
    C.encabezar(ws, 'Riesgos de la línea de fritura',
                'Los riesgos de la línea con su probabilidad, su gravedad y su medida. No '
                'sustituye a la evaluación de riesgos de tu servicio de prevención: es la lista '
                'de lo que tiene que contemplar en la fritura.', col_fin='K')
    refs, fila_listas = C.bloque_listas(
        ws, P['listas'], [('Sí o No', [SI, NO]), ('Escala de 1 a 3', [1, 2, 3])], col='A')
    motor.val(ws, 'B%d' % P['alto'], 'Nivel desde el que la prioridad es ALTA (probabilidad por '
                                     'gravedad)', wrap=True, bold=True)
    entrada(ws, 'C%d' % P['alto'], 6, fmt=C.FMT_ENT, etiqueta='Umbral de prioridad alta',
            align='center')
    C.nota(ws, 'D%d' % P['alto'], 'Supuesto de la casa: con una escala de 1 a 3, un 6 o un 9 '
                                  'piden medida antes de abrir.')
    motor.val(ws, 'B%d' % P['medio'], 'Nivel desde el que la prioridad es MEDIA', bold=True)
    entrada(ws, 'C%d' % P['medio'], 3, fmt=C.FMT_ENT, etiqueta='Umbral de prioridad media',
            align='center')
    C.cabecera(ws, P['cab'], [('A', 'Nº'), ('B', 'Riesgo'), ('C', 'Dónde'), ('D', 'Medida'),
                              ('E', 'Probabilidad (1-3)'), ('F', 'Gravedad (1-3)'),
                              ('G', 'Nivel'), ('H', 'Prioridad'), ('I', '¿Medida puesta?'),
                              ('J', 'Responsable'), ('K', 'Fuente')])
    escalas, medidas = [], []
    for i, (riesgo, donde, medida, prob, grav, fuente) in enumerate(RIESGOS):
        r = P['ini'] + i
        motor.val(ws, 'A%d' % r, i + 1, fmt=C.FMT_ENT, align='center')
        motor.val(ws, 'B%d' % r, riesgo, wrap=True)
        motor.val(ws, 'C%d' % r, donde, wrap=True)
        motor.val(ws, 'D%d' % r, medida, wrap=True)
        entrada(ws, 'E%d' % r, prob, fmt=C.FMT_ENT, etiqueta='Probabilidad: ' + riesgo,
                align='center')
        entrada(ws, 'F%d' % r, grav, fmt=C.FMT_ENT, etiqueta='Gravedad: ' + riesgo, align='center')
        escalas += ['E%d' % r, 'F%d' % r]
        motor.f(ws, 'G%d' % r, ie('E%d*F%d' % (r, r)), fmt=C.FMT_ENT, align='center', bold=True)
        motor.f(ws, 'H%d' % r,
                ie('IF(G{r}>=$C${a},"Alta",IF(G{r}>=$C${m},"Media","Baja"))'
                   .format(r=r, a=P['alto'], m=P['medio'])), align='center', bold=True)
        entrada(ws, 'I%d' % r, NO, etiqueta='Medida puesta: ' + riesgo, align='center')
        medidas.append('I%d' % r)
        entrada(ws, 'J%d' % r, 'Titular', etiqueta='Responsable: ' + riesgo)
        motor.val(ws, 'K%d' % r, '[Fuente: %s]' % (fuente or 'supuesto (buena práctica de la casa)'),
                  wrap=True)
        ws.row_dimensions[r].height = 40
        if fuente:
            notas_de_fuente(ws, ['D%d' % r, 'B%d' % r], fuente)
    C.dv_rango(ws, escalas, refs['Escala de 1 a 3'], 'De 1 a 3',
               '1 = baja, 2 = media, 3 = alta.')
    C.dv_rango(ws, medidas, refs['Sí o No'], 'Contesta Sí o No', 'Elige «Sí» o «No».')
    motor.semaforo_texto(ws, 'H%d:H%d' % (P['ini'], P['fin']),
                         (('Alta', motor.CF_ROJO_BG, motor.CF_ROJO_FG),
                          ('Media', motor.CF_AMBAR_BG, motor.CF_AMBAR_FG),
                          ('Baja', motor.CF_VERDE_BG, motor.CF_VERDE_FG)))

    def res(fila, etiqueta, formula):
        motor.val(ws, 'B%d' % fila, etiqueta, bold=True)
        motor.f(ws, 'C%d' % fila, formula, fmt=C.FMT_ENT, bold=True, align='center')

    res(P['n_alta'], 'Riesgos de prioridad alta',
        ie('COUNTIF(H%d:H%d,"Alta")' % (P['ini'], P['fin'])))
    res(P['n_media'], 'Riesgos de prioridad media',
        ie('COUNTIF(H%d:H%d,"Media")' % (P['ini'], P['fin'])))
    res(P['pend_alta'], 'De prioridad alta, sin la medida puesta',
        '=SUMPRODUCT(--(H{a}:H{b}="Alta"),--(I{a}:I{b}="{n}"))'.format(a=P['ini'], b=P['fin'], n=NO))
    motor.semaforo_isnumber(ws, 'C%d' % P['pend_alta'], '$C$%d' % P['pend_alta'], '>', '0')
    motor.val(ws, 'B%d' % P['ver'], 'VEREDICTO', bold=True)
    motor.f(ws, 'C%d' % P['ver'],
            ie('IF(C{0}>0,"No abras la línea con riesgos altos sin su medida: te faltan "&C{0},'
               '"Los riesgos altos tienen su medida puesta")'.format(P['pend_alta'])), bold=True)
    ws.merge_cells('C%d:K%d' % (P['ver'], P['ver']))
    C.destacado(ws, 'C%d' % P['ver'])
    C.parrafo(ws, P['nota'], 'Las probabilidades y gravedades sembradas son SUPUESTOS de la casa '
              'para El Molinete: cámbialas con tu servicio de prevención. La nocturnidad y los '
              'descansos de la plantilla se calculan en turnos-plantilla-y-madrugada.xlsx.', 'A',
              'K', alto=34)
    C.parrafo(ws, P['nota'] + 1, 'Nunca agua sobre aceite ardiendo: el extintor de la freidora '
              'es de clase F, y la tapa de la cuba, a mano.', 'A', 'K', alto=24)
    ws.freeze_panes = 'C%d' % (P['cab'] + 1)
    C.pagina(ws, titulos='$%d:$%d' % (P['cab'], P['cab']), area='A1:K%d' % fila_listas)
    return ws


# ==========================================================================
# Hoja «Registro de Formación»
# ==========================================================================
def hoja_formacion(wb):
    ws = wb.create_sheet(H_FOR)
    C.anchos(ws, {'A': 6, 'B': 34, 'C': 28, 'D': 12, 'E': 46, 'F': 16, 'G': 16, 'H': 18,
                  'I': 60})
    C.encabezar(ws, 'Registro de formación del personal',
                'Lo que la norma exige es que la EMPRESA garantice la formación de cada '
                'persona de acuerdo con su puesto y pueda ACREDITARLA. Eso es esta tabla.',
                col_fin='I')
    C.seccion(ws, 'A5', 'Una línea por persona')
    C.cabecera(ws, F_CAB, [('A', 'Nº'), ('B', 'Persona'), ('C', 'Puesto'), ('D', 'Jornada'),
                           ('E', 'Contenido de la formación'), ('F', 'Fecha de la formación'),
                           ('G', 'Caduca el'), ('H', 'Estado'), ('I', 'Notas')])
    v_fecha = []
    for i in range(F_FIN - F_INI + 1):
        r = F_INI + i
        if i < len(D.PLANTILLA):
            p = D.PLANTILLA[i]
            persona, puesto = 'Persona %d' % (i + 1), p['puesto']
            jornada = p['jornada']
            if p['es_titular']:
                contenido = ('Responsable del autocontrol: higiene, alérgenos, aceite de fritura '
                             'y registros')
            elif 'Churrero' in puesto:
                contenido = 'Higiene, alérgenos, aceite de fritura y extintor, en la línea'
            elif 'Camarero' in puesto:
                contenido = 'Higiene, alérgenos e información al cliente, en la barra y la sala'
            else:
                contenido = 'Higiene, alérgenos y aceite de fritura, en la línea y la barra'
        else:
            persona, puesto, jornada, contenido = LIBRE_P, 'Pendiente', 1.0, 'Pendiente'
        motor.val(ws, 'A%d' % r, i + 1, fmt=C.FMT_ENT, align='center')
        entrada(ws, 'B%d' % r, persona, etiqueta='Persona %d' % (i + 1))
        entrada(ws, 'C%d' % r, puesto, etiqueta='Puesto %d' % (i + 1))
        entrada(ws, 'D%d' % r, jornada, fmt='0.0', etiqueta='Jornada %d' % (i + 1), align='center')
        entrada(ws, 'E%d' % r, contenido, etiqueta='Contenido %d' % (i + 1), wrap=True)
        entrada(ws, 'F%d' % r, FECHA_BASE, fmt=C.FMT_FECHA, etiqueta='Fecha %d' % (i + 1))
        v_fecha.append('F%d' % r)
        motor.f(ws, 'G%d' % r,
                ie('IF($B{r}="{lb}","",DATE(YEAR($F{r}),MONTH($F{r})+$C${v},DAY($F{r})))'
                   .format(r=r, lb=LIBRE_P, v=F_VALIDEZ)), fmt=C.FMT_FECHA)
        motor.f(ws, 'H%d' % r,
                ie('IF($B{r}="{lb}","",IF($G{r}<$C${c},"CADUCADA",IF($G{r}<$C${c}+$C${p},'
                   '"Caduca pronto","Vigente")))'.format(r=r, lb=LIBRE_P, c=F_CONTROL, p=F_AVISO)),
                bold=True)
        ws.row_dimensions[r].height = 30
    nota_celda(ws, 'E%d' % F_CAB, 'CHN-69')
    motor.dv_fecha(ws, v_fecha)
    motor.semaforo_texto(ws, 'H%d:H%d' % (F_INI, F_FIN),
                         (('CADUCADA', motor.CF_ROJO_BG, motor.CF_ROJO_FG),
                          ('Caduca pronto', motor.CF_AMBAR_BG, motor.CF_AMBAR_FG),
                          ('Vigente', motor.CF_VERDE_BG, motor.CF_VERDE_FG)))
    C.nota(ws, 'I%d' % F_INI, 'Las cinco personas de El Molinete (titular, dos churreros/as, '
                              'camarero/a y refuerzo). Las filas libres son para tu equipo.')

    C.seccion(ws, 'A%d' % F_SEC_PAR, 'Parámetros de esta hoja')
    for fila, etiq, valor, fmt, nota_txt in (
            (F_VALIDEZ, 'Validez que TÚ fijas (meses)', 24, C.FMT_ENT,
             'La norma no fija una caducidad: lo que hay es una obligación continua del '
             'empresario. Dos años es un criterio de la casa para que el registro no se quede '
             'dormido; déjalo escrito en tu plan de autocontrol.'),
            (F_CONTROL, 'Fecha de control', FECHA_BASE, C.FMT_FECHA,
             'La fecha desde la que se mira si algo ha caducado. Ponla en «hoy» cada vez que '
             'revises el registro.'),
            (F_AVISO, 'Aviso de «caduca pronto» (días antes)', 60, C.FMT_ENT,
             'Supuesto: dos meses dan tiempo a organizar una sesión sin parar la línea.')):
        motor.val(ws, 'A%d' % fila, etiq, bold=True)
        ws.merge_cells('A%d:B%d' % (fila, fila))
        entrada(ws, 'C%d' % fila, valor, fmt=fmt, etiqueta=etiq, align='center')
        C.nota(ws, 'I%d' % fila, nota_txt)
    motor.dv_fecha(ws, ['C%d' % F_CONTROL])

    C.seccion(ws, 'A%d' % F_SEC_RES, 'Resumen')

    def res(fila, etiqueta, formula, fmt=None):
        motor.val(ws, 'B%d' % fila, etiqueta, bold=True, wrap=True)
        motor.f(ws, 'C%d' % fila, formula, fmt=fmt, bold=True)

    res(F_CON, 'Personas con formación registrada',
        '=SUMPRODUCT(--($B${a}:$B${b}<>"{lb}"))'.format(a=F_INI, b=F_FIN, lb=LIBRE_P), C.FMT_ENT)
    res(F_CADUCADAS, 'Formaciones CADUCADAS', ie('COUNTIF(H%d:H%d,"CADUCADA")' % (F_INI, F_FIN)),
        C.FMT_ENT)
    res(F_PROXIMAS, 'Formaciones que caducan pronto',
        ie('COUNTIF(H%d:H%d,"Caduca pronto")' % (F_INI, F_FIN)), C.FMT_ENT)
    res(F_VER, 'VEREDICTO',
        ie('IF(C{c}>0,"Tienes formación caducada: es lo primero que se pide en una inspección",'
           'IF(C{p}>0,"Programa ya las que caducan pronto","Registro al día"))'
           .format(c=F_CADUCADAS, p=F_PROXIMAS)))
    ws.merge_cells('C%d:H%d' % (F_VER, F_VER))
    C.destacado(ws, 'C%d' % F_VER)
    motor.semaforo_isnumber(ws, 'C%d' % F_CADUCADAS, '$C$%d' % F_CADUCADAS, '>', '0')
    motor.semaforo_isnumber(ws, 'C%d' % F_PROXIMAS, '$C$%d' % F_PROXIMAS, '>', '0',
                            bg=motor.CF_AMBAR_BG, fg=motor.CF_AMBAR_FG)
    C.parrafo(ws, F_NOTA, 'Lo que vale ante una inspección es este registro, con el contenido '
              'y la fecha de cada formación, no un papel suelto con un sello. La '
              'responsabilidad de tenerlo es de la empresa.', 'A', 'I', alto=34)
    C.pie(ws, F_PIE, 'A', 'E')
    ws.freeze_panes = 'B%d' % F_INI
    C.pagina(ws, titulos='$%d:$%d' % (F_CAB, F_CAB), area='A1:I%d' % (F_PIE + 2))
    return ws


# ==========================================================================
# Hoja «Cronograma y Ruta Crítica»
# ==========================================================================
NOTAS_GANTT = {
    'H3': 'El proyecto de la extracción es el que decide si hay licencia de obra (1.ª fila del '
          'checklist).',
    'H4': 'Sólo si la primera fila del checklist dijo que sí. Si no hay obra con proyecto, pon '
          'su duración a cero.',
    'H7': 'SU DURACIÓN SALE DEL PLAZO DE ENTREGA (celda C6, en semanas): es la fuente única del '
          'paquete. Cámbialo por el que te den por escrito.',
    'H10': 'Comunicación o declaración responsable a tu comunidad: preséntala antes de vender.',
    'H13': 'Las pruebas de masa no son un trámite: son la semana en que descubres si tu línea y '
           'tu receta se llevan bien.',
}


def hoja_gantt(wb):
    ws = wb.create_sheet(H_GAN)
    C.anchos(ws, dict([('A', 8), ('B', 50)] + [(c, 10) for c in COLS_MAT]
                      + [('Q', 4), ('R', 80)]))
    C.encabezar(ws, 'Cronograma y ruta crítica de la apertura',
                'Catorce hitos con su duración y sus dependencias, en MESES y con aritmética de '
                'índices. El plazo de entrega de la maquinaria crítica se pone AQUÍ, y sólo aquí.',
                col_fin='R')
    refs, fila_listas = C.bloque_listas(
        ws, GN['listas'],
        [('Hitos y «sin dependencia»', [SIN_DEP] + [h[0] for h in D.GANTT]),
         ('Meses del año', list(D.MESES))], col='A')
    ref_meses = refs['Meses del año'].lstrip('=')

    def param(clave, etiqueta, valor, fmt, nota_txt, verde=True):
        motor.val(ws, 'A%d' % GN[clave], etiqueta, bold=True)
        ws.merge_cells('A%d:B%d' % (GN[clave], GN[clave]))
        if verde:
            entrada(ws, 'C%d' % GN[clave], valor, fmt=fmt, etiqueta=etiqueta, align='center')
        else:
            motor.val(ws, 'C%d' % GN[clave], valor, fmt=fmt, align='center')
        C.nota(ws, 'R%d' % GN[clave], nota_txt)

    param('arranque', 'Mes en que arranca el proyecto (1-12)', D.MES_INICIO_PROYECTO, C.FMT_ENT,
          'El mes natural del primer paso. El ejemplo arranca en noviembre para abrir el 1 de '
          'septiembre, un mes antes del pico de octubre a marzo.')
    semanas, _f, nota_plazo = D.PLAZO_ENTREGA_MAQUINARIA_SEMANAS
    param('plazo', 'Plazo de entrega de la maquinaria crítica (semanas)', semanas, C.FMT_ENT,
          'FUENTE ÚNICA DEL PAQUETE: ningún otro libro lo pide. Supuesto declarado de El '
          'Molinete: ningún distribuidor publica plazos. ' + nota_plazo)
    param('dias_sem', 'Días de una semana', 7, C.FMT_ENT,
          'Constante de conversión en celda, para que no viva dentro de ninguna fórmula.')
    param('dias_mes', 'Días de un mes medio', 365.0 / 12.0, C.FMT_DEC,
          'Un año de 365 días repartido en 12 meses.')
    motor.val(ws, 'A%d' % GN['sem_mes'], 'Semanas de un mes medio', bold=True)
    ws.merge_cells('A%d:B%d' % (GN['sem_mes'], GN['sem_mes']))
    motor.f(ws, 'C%d' % GN['sem_mes'], ie('C%d/C%d' % (GN['dias_mes'], GN['dias_sem'])),
            fmt=C.FMT_DEC, align='center')
    param('centinela', 'Valor centinela de la matriz de dependencias (meses)', 999, C.FMT_ENT,
          'Truco de cálculo, no un dato: las casillas de la matriz que no aplican llevan este '
          'número para que MIN trabaje sólo con números. Por eso NO está en verde.', verde=False)
    param('objetivo', 'Mes en que quieres abrir (1-12)', D.dato(D.NEGOCIO, 'mes_apertura'),
          C.FMT_ENT, 'Tu objetivo. El de El Molinete es septiembre: rodaje antes del pico.')

    C.cabecera(ws, GN['cab'],
               [('A', 'Id'), ('B', 'Hito'), ('C', 'Duración (meses)'), ('D', 'Depende de (1)'),
                ('E', 'Depende de (2)'), ('F', 'Depende de (3)'),
                ('G', 'Empieza (mes del proyecto)'), ('H', 'Acaba (mes del proyecto)'),
                ('I', 'Holgura (meses)'), ('J', '¿Ruta crítica?'), ('K', 'Mes natural de fin'),
                ('R', 'Notas')])
    deps_coords = []
    for i, (hid, hito, dur, deps) in enumerate(D.GANTT):
        fila = GN['ini'] + i
        motor.val(ws, 'A%d' % fila, hid, align='center')
        motor.val(ws, 'B%d' % fila, hito, wrap=True)
        if dur is None:
            motor.f(ws, 'C%d' % fila, ie('$C$%d/$C$%d' % (GN['plazo'], GN['sem_mes'])),
                    fmt=C.FMT_DEC1, align='center', bold=True)
        else:
            entrada(ws, 'C%d' % fila, dur, fmt=C.FMT_DEC1, etiqueta='Duración de ' + hid,
                    align='center')
        lista = list(deps) + [SIN_DEP] * (3 - len(deps))
        for j, letra in enumerate(('D', 'E', 'F')):
            entrada(ws, '%s%d' % (letra, fila), lista[j], etiqueta='Dependencia %d de %s'
                    % (j + 1, hid), align='center')
            deps_coords.append('%s%d' % (letra, fila))
        arriba = fila - 1
        trozos = ['IFERROR(INDEX($H$%d:H%d,MATCH(%s%d,$A$%d:A%d,0)),0)'
                  % (GN['cab'], arriba, letra, fila, GN['cab'], arriba) for letra in 'DEF']
        motor.f(ws, 'G%d' % fila, ie('MAX(%s)' % ','.join(trozos)), fmt=C.FMT_DEC1)
        motor.f(ws, 'H%d' % fila, ie('G%d+C%d' % (fila, fila)), fmt=C.FMT_DEC1)
        if hid in NOTAS_GANTT:
            C.nota(ws, 'R%d' % fila, NOTAS_GANTT[hid])
        ws.row_dimensions[fila].height = 26
    C.dv_rango(ws, deps_coords, refs['Hitos y «sin dependencia»'], 'Elige un hito o «-»',
               'Identificador de otro hito MÁS ARRIBA en la tabla, o «-» si no depende de ninguno.')
    nota_celda(ws, 'B%d' % (GN['ini'] + 3), 'CUN-35')

    barra_seccion(ws, GN['sec_mat'], 'MATRIZ DE DEPENDENCIAS: QUÉ HITO BLOQUEA A CUÁL',
                  ['B'] + list(COLS_MAT) + ['R'])
    C.cabecera(ws, GN['mat_cab'], [('A', 'Sucesor'), ('B', 'empieza en el mes...')]
               + [(COLS_MAT[k], D.GANTT[k][0]) for k in range(len(D.GANTT))] + [('R', 'Notas')],
               altura=22)
    for i, (hid, _h, _d, _deps) in enumerate(D.GANTT):
        fila, fila_h = GN['mat_ini'] + i, GN['ini'] + i
        motor.val(ws, 'A%d' % fila, hid, align='center')
        motor.f(ws, 'B%d' % fila, ie('G%d' % fila_h), fmt=C.FMT_DEC1)
        for k in range(len(D.GANTT)):
            fila_pred = GN['ini'] + k
            motor.f(ws, '%s%d' % (COLS_MAT[k], fila),
                    ie('IF(OR($D${0}=$A${1},$E${0}=$A${1},$F${0}=$A${1}),$G${0},$C${2})'
                       .format(fila_h, fila_pred, GN['centinela'])), fmt=FMT_CENTINELA)
    C.nota(ws, 'R%d' % GN['mat_ini'], 'Se lee por COLUMNAS: en la columna de H3 aparece el mes '
                                      'en que empieza cada hito que depende de H3. Las casillas '
                                      'que parecen vacías llevan el centinela.')
    for i, _h in enumerate(D.GANTT):
        fila_h, col = GN['ini'] + i, COLS_MAT[i]
        rango = '%s%d:%s%d' % (col, GN['mat_ini'], col, GN['mat_fin'])
        motor.f(ws, 'I%d' % fila_h, ie('MIN(MIN({0}),MAX($H${1}:$H${2}))-H{3}'
                                       .format(rango, GN['ini'], GN['fin'], fila_h)),
                fmt=C.FMT_DEC1)
        motor.f(ws, 'J%d' % fila_h, ie('IF(ROUND(I{0},3)<=0,"{1}","{2}")'.format(fila_h, SI, NO)))
        motor.f(ws, 'K%d' % fila_h, ie('INDEX({0},ROUNDDOWN(MOD($C${1}-1+H{2},12),0)+1)'
                                       .format(ref_meses, GN['arranque'], fila_h)))
    motor.regla_expresion(ws, 'J%d:J%d' % (GN['ini'], GN['fin']), '=$J%d="%s"' % (GN['ini'], SI),
                          bg=motor.CF_AMBAR_BG, fg=motor.CF_AMBAR_FG)

    barra_seccion(ws, GN['sec_res'], 'LO QUE SALE DE TODO ESTO', ['B'] + list(COLS_MAT) + ['R'])

    def res(clave, etiqueta, formula, fmt=None, nota_txt=None, bold=True, ancho=False):
        fila = GN[clave]
        motor.val(ws, 'A%d' % fila, etiqueta, wrap=True, bold=bold)
        ws.merge_cells('A%d:B%d' % (fila, fila))
        motor.f(ws, 'C%d' % fila, formula, fmt=fmt, bold=bold)
        if ancho:
            ws.merge_cells('C%d:P%d' % (fila, fila))
            ws['C%d' % fila].alignment = Alignment(vertical='top', wrap_text=True)
            ws.row_dimensions[fila].height = 34
        if nota_txt:
            C.nota(ws, 'R%d' % fila, nota_txt)

    res('duracion', 'Duración total del proyecto (meses)', ie('MAX(H%d:H%d)' % (GN['ini'], GN['fin'])),
        C.FMT_DEC1, 'Desde el primer día de búsqueda de local hasta el día que abres.')
    res('criticos', 'Hitos SIN holgura (ruta crítica)',
        ie('COUNTIF(J%d:J%d,"%s")' % (GN['ini'], GN['fin'], SI)), C.FMT_ENT,
        'Si uno de éstos se retrasa un día, la apertura se retrasa un día.')
    res('mes_ap', 'Mes natural en el que abres (1-12)',
        ie('ROUNDDOWN(MOD($C${0}-1+C{1},12),0)+1'.format(GN['arranque'], GN['duracion'])),
        C.FMT_ENT)
    res('nombre_mes', 'Nombre de ese mes', ie('INDEX(%s,C%d)' % (ref_meses, GN['mes_ap'])))
    res('disponibles', 'Meses que tienes entre el arranque y tu objetivo',
        ie('IF(MOD(C{0}-C{1},12)=0,12,MOD(C{0}-C{1},12))'.format(GN['objetivo'], GN['arranque'])),
        C.FMT_DEC1,
        'Si el objetivo cae en el mismo mes del arranque, la hoja entiende que es el del año '
        'siguiente (doce meses).')
    res('margen', 'Margen entre lo que tienes y lo que necesitas (meses)',
        ie('C{0}-C{1}'.format(GN['disponibles'], GN['duracion'])), C.FMT_DEC1,
        'Negativo = no llegas a tu objetivo con esta ruta crítica.')
    res('llegas', '¿Llegas a abrir en tu objetivo?',
        ie('IF(ROUND(C{0},2)<0,"NO LLEGAS: arranca antes, acorta la ruta crítica o mueve el '
           'objetivo",IF(ROUND(C{0},2)=0,"LLEGAS JUSTO: cualquier retraso en la ruta crítica te '
           'mueve la apertura","LLEGAS CON MARGEN"))'.format(GN['margen'])), ancho=True)
    motor.semaforo_texto(ws, 'C%d' % GN['llegas'],
                         (('NO LLEGAS: arranca antes, acorta la ruta crítica o mueve el objetivo',
                           motor.CF_ROJO_BG, motor.CF_ROJO_FG),
                          ('LLEGAS JUSTO: cualquier retraso en la ruta crítica te mueve la '
                           'apertura', motor.CF_AMBAR_BG, motor.CF_AMBAR_FG),
                          ('LLEGAS CON MARGEN', motor.CF_VERDE_BG, motor.CF_VERDE_FG)))
    C.destacado(ws, 'C%d' % GN['llegas'])
    res('maq_critica', '¿La maquinaria está en la ruta crítica?', ie('J%d' % FILA_H7),
        nota_txt='Si dice «Sí», tu fecha de apertura la manda el proveedor, no el papeleo.')
    res('maq_holgura', 'Holgura de la maquinaria (meses)', ie('I%d' % FILA_H7), C.FMT_DEC1,
        'Lo que puede retrasarse la entrega sin mover la apertura.')
    motor.semaforo_texto(ws, 'C%d' % GN['maq_critica'], ((SI, motor.CF_AMBAR_BG, motor.CF_AMBAR_FG),))

    C.parrafo(ws, GN['nota'], 'LAS DURACIONES SON SUPUESTOS y casi todas dependen de un '
              'tercero. Cámbialas por las tuyas en cuanto tengas un plazo por escrito, empezando '
              'por el de la maquinaria: es el que más silenciosamente mueve la fecha, porque no '
              'falla, sólo llega tarde.', 'A', 'R', alto=40)
    C.parrafo(ws, GN['nota'] + 1, 'La hoja no usa funciones de fecha a propósito: trabaja en '
              'meses decimales, así que medio mes es 0,5 y se suma sin discutir con el '
              'calendario laboral.', 'A', 'R', alto=30)
    ws.freeze_panes = 'C%d' % (GN['cab'] + 1)
    C.pagina(ws, titulos='$%d:$%d' % (GN['cab'], GN['cab']), area='A1:R%d' % fila_listas)
    return ws


# ==========================================================================
# Mapa de celdas citables
# ==========================================================================
def mapa_celdas():
    m = [
        ('¿La extracción necesita proyecto? (1.ª fila del checklist)', H_CHK, 'K%d' % K_EXTRACCION, 'entrada'),
        ('Lo que te toca por la extracción', H_CHK, 'L%d' % K_EXTRACCION, 'salida'),
        ('¿Tu actividad está en el Anexo de la Ley 12/2012? (2.ª fila)', H_CHK, 'K%d' % K_ANEXO, 'salida'),
        ('Lo que te toca por el Anexo', H_CHK, 'L%d' % K_ANEXO, 'salida'),
        ('Trámites del expediente', H_CHK, 'D%d' % K_TOTAL, 'salida'),
        ('Trámites hechos', H_CHK, 'D%d' % K_HECHOS, 'salida'),
        ('Avance del expediente', H_CHK, 'D%d' % K_AVANCE, 'salida'),
        ('Trámites sin importe (a presupuestar)', H_CHK, 'D%d' % K_SINIMP, 'salida'),
        ('Trámites que cambian según la comunidad autónoma', H_CHK, 'D%d' % K_CCAA_N, 'salida'),
        ('Avance de la fase F1', H_CHK, 'J%d' % FASES[0][5], 'salida'),
        ('Trámites de la fase F1', H_CHK, 'E%d' % FASES[0][5], 'salida'),
        ('Plazo en días del pedido de maquinaria crítica', H_CHK, 'E%d' % K_MAQ, 'salida'),
        ('Comunidad autónoma del ejemplo', H_CHK, 'D%d' % K_TU_CCAA, 'entrada'),
        ('Lo que vale para ti de los ejemplos de Madrid', H_CHK, 'D%d' % K_TU_LECT, 'salida'),
        ('¿Consumo en el local?', H_IAE, 'B%d' % I['consumo'], 'entrada'),
        ('¿Venta para llevar?', H_IAE, 'B%d' % I['llevar'], 'entrada'),
        ('¿Ferias o casetas?', H_IAE, 'B%d' % I['ferias'], 'entrada'),
        ('¿Sirves a otros negocios?', H_IAE, 'B%d' % I['otros'], 'entrada'),
        ('Epígrafe de la sala elegido', H_IAE, 'B%d' % I['sala'], 'entrada'),
        ('Formatos que te tocan', H_IAE, 'B%d' % I['n'], 'salida'),
        ('Epígrafe candidato de la sala', H_IAE, 'C%d' % I['ini'], 'salida'),
        ('¿Necesitas el 644.6 además de la sala?', H_IAE, 'B%d' % I['p6446'], 'salida'),
        ('Epígrafes candidatos que llevas a tu asesor', H_IAE, 'B%d' % I['epigrafes'], 'salida'),
        ('¿Pagas el impuesto? (exención)', H_IAE, 'B%d' % I['pago'], 'salida'),
        ('La pregunta del CNAE para tu asesor', H_IAE, 'B%d' % I['pregunta'], 'salida'),
        ('Tu actividad: régimen de apertura', H_REG, 'B%d' % R['actividad'], 'salida'),
        ('Tus obras: licencia aparte', H_REG, 'B%d' % R['obras'], 'salida'),
        ('Tu terraza: autorización aparte', H_REG, 'B%d' % R['terraza_v'], 'salida'),
        ('Clasificación horaria del ejemplo', H_REG, 'B%d' % R['clas'], 'entrada'),
        ('Hora mínima de apertura de cafetería o bar en Madrid', H_REG, 'B%d' % R['cafe'], 'parametro'),
        ('Hora mínima de apertura de chocolatería en Madrid', H_REG, 'B%d' % R['choco'], 'parametro'),
        ('Hora mínima de apertura de tu clasificación', H_REG, 'B%d' % R['hmin'], 'salida'),
        ('Lo que significa el horario para tu desayuno', H_REG, 'B%d' % R['lectura'], 'salida'),
        ('Formato para humos y gas', H_HUM, 'B%d' % G['formato'], 'entrada'),
        ('Gas o eléctrica', H_HUM, 'B%d' % G['energia'], 'entrada'),
        ('Años entre inspecciones del gas', H_HUM, 'B%d' % G['anios_gas'], 'parametro'),
        ('Peso de bombona sin instalación receptora (kg)', H_HUM, 'B%d' % G['kg_bombona'], 'parametro'),
        ('Licencia de obra', H_HUM, 'B%d' % G['obra'], 'salida'),
        ('Salida de humos del ejemplo de Madrid', H_HUM, 'B%d' % G['humos'], 'salida'),
        ('Gas: lo que te toca', H_HUM, 'B%d' % G['gas'], 'salida'),
        ('Próxima inspección del gas', H_HUM, 'B%d' % G['prox_gas'], 'salida'),
        ('Extintor de clase F', H_HUM, 'B%d' % G['ext_v'], 'salida'),
        ('Papeles de obra, humos y gas que te tocan', H_HUM, 'B%d' % G['n_toca'], 'salida'),
        ('Veredicto de obra, humos y gas', H_HUM, 'B%d' % G['ver'], 'salida'),
        ('Tu régimen sanitario', H_SAN, 'B%d' % S['regimen'], 'salida'),
        ('El trámite sanitario que te toca', H_SAN, 'B%d' % S['tramite'], 'salida'),
        ('Estado de la autorización de feria', H_FER, 'B%d' % FE['estado'], 'salida'),
        ('Requisitos del puesto que te aplican', H_FER, 'B%d' % FE['aplican'], 'salida'),
        ('Veredicto de ferias', H_FER, 'B%d' % FE['ver'], 'salida'),
        ('¿Freidora dedicada o compartida?', H_ALE, 'B%d' % A['dedicada'], 'salida'),
        ('Casillas de alérgenos pendientes de la etiqueta', H_ALE, 'B%d' % A['revisar'], 'salida'),
        ('¿Declaras el gluten por escrito?', H_ALE, 'B%d' % A['declara'], 'salida'),
        ('Umbral del «sin gluten» (mg/kg)', H_ALE, 'B%d' % A['sg_umbral'], 'parametro'),
        ('Acrilamida: parte A', H_ALE, 'B%d' % A['ac_a'], 'salida'),
        ('Acrilamida: parte B', H_ALE, 'B%d' % A['ac_b'], 'salida'),
        ('Riesgos de prioridad alta', H_PRL, 'C%d' % P['n_alta'], 'salida'),
        ('Riesgos altos sin medida', H_PRL, 'C%d' % P['pend_alta'], 'salida'),
        ('Veredicto de riesgos de la fritura', H_PRL, 'C%d' % P['ver'], 'salida'),
        ('Personas con formación registrada', H_FOR, 'C%d' % F_CON, 'salida'),
        ('Validez de la formación en meses', H_FOR, 'C%d' % F_VALIDEZ, 'parametro'),
        ('Veredicto del registro de formación', H_FOR, 'C%d' % F_VER, 'salida'),
        ('Mes en que arranca el proyecto', H_GAN, 'C%d' % GN['arranque'], 'entrada'),
        ('Plazo de entrega de la maquinaria crítica (semanas)', H_GAN, 'C%d' % GN['plazo'], 'entrada'),
        ('Duración del pedido de maquinaria (meses)', H_GAN, 'C%d' % FILA_H7, 'salida'),
        ('Duración total del proyecto (meses)', H_GAN, 'C%d' % GN['duracion'], 'salida'),
        ('Hitos sin holgura (ruta crítica)', H_GAN, 'C%d' % GN['criticos'], 'salida'),
        ('Mes natural en el que abres', H_GAN, 'C%d' % GN['mes_ap'], 'salida'),
        ('Nombre del mes de apertura', H_GAN, 'C%d' % GN['nombre_mes'], 'salida'),
        ('Mes en que quieres abrir', H_GAN, 'C%d' % GN['objetivo'], 'entrada'),
        ('¿Llegas a abrir en tu objetivo?', H_GAN, 'C%d' % GN['llegas'], 'salida'),
        ('¿La maquinaria está en la ruta crítica?', H_GAN, 'C%d' % GN['maq_critica'], 'salida'),
        ('Holgura de la maquinaria (meses)', H_GAN, 'C%d' % GN['maq_holgura'], 'salida'),
    ]
    return m


# ==========================================================================
# Demostraciones con pycel
# ==========================================================================
def demo(ruta):
    from pycel import ExcelCompiler
    ok, fallos = [], []

    def prueba(nombre, cond, detalle=''):
        (ok if cond else fallos).append(nombre + (' - ' + detalle if detalle else ''))

    def ev(xl, hoja, coord):
        return xl.evaluate("'%s'!%s" % (hoja, coord))

    exc = ExcelCompiler(ruta)
    # 1. Ruta crítica y mes de apertura = juego de datos.
    total = ev(exc, H_GAN, 'C%d' % GN['duracion'])
    esperado, _c, holg = D.ruta_critica()
    prueba('la duración total coincide con la ruta crítica de datos_ejemplo',
           abs(total - esperado) < 0.001, '%r frente a %r meses' % (total, esperado))
    mes = ev(exc, H_GAN, 'C%d' % GN['mes_ap'])
    prueba('el mes de apertura es el del juego de datos (septiembre)',
           int(mes) == D.mes_apertura_calculado() == D.dato(D.NEGOCIO, 'mes_apertura'),
           '%r' % mes)
    hm = ev(exc, H_GAN, 'C%d' % GN['maq_holgura'])
    prueba('la holgura de la maquinaria es la de datos_ejemplo', abs(hm - holg['H7']) < 0.01,
           '%r frente a %r' % (hm, holg['H7']))

    # 2. Cambiar el plazo de la maquinaria mueve la apertura (fuente única).
    x2 = ExcelCompiler(ruta)
    for c in ('duracion', 'llegas', 'maq_critica'):
        ev(x2, H_GAN, 'C%d' % GN[c])
    ev(x2, H_CHK, 'E%d' % K_MAQ)
    x2.set_value("'%s'!C%d" % (H_GAN, GN['plazo']), 30)
    dur2 = ev(x2, H_GAN, 'C%d' % GN['duracion'])
    prueba('con 30 semanas de entrega, la maquinaria entra en la ruta crítica y no llegas',
           dur2 > esperado and ev(x2, H_GAN, 'C%d' % GN['maq_critica']) == SI
           and str(ev(x2, H_GAN, 'C%d' % GN['llegas'])).startswith('NO LLEGAS'),
           '%r meses' % dur2)
    prueba('la fila del checklist que pide la maquinaria lee el mismo plazo',
           ev(x2, H_CHK, 'E%d' % K_MAQ) == 30 * 7, '%r días' % ev(x2, H_CHK, 'E%d' % K_MAQ))

    # 3. D13: la 1.ª fila; y sin mesas, la actividad entra en el Anexo.
    l1 = ev(exc, H_CHK, 'L%d' % K_EXTRACCION)
    prueba('la 1.ª fila manda a licencia de obra y técnico', 'Licencia de obra y técnico' in l1, l1[:60])
    prueba('con sala, la 2.ª fila dice que la actividad sale del Anexo',
           ev(exc, H_CHK, 'K%d' % K_ANEXO) == NO)
    x3 = ExcelCompiler(ruta)
    ev(x3, H_CHK, 'K%d' % K_ANEXO)
    ev(x3, H_REG, 'B%d' % R['actividad'])
    ev(x3, H_IAE, 'B%d' % I['pregunta'])
    x3.set_value("'%s'!B%d" % (H_IAE, I['consumo']), NO)
    act = ev(x3, H_REG, 'B%d' % R['actividad'])
    prueba('sólo despacho: Anexo = Sí y la actividad va por declaración responsable',
           ev(x3, H_CHK, 'K%d' % K_ANEXO) == SI and 'declaración responsable' in act, act[:70])
    preg = ev(x3, H_IAE, 'B%d' % I['pregunta'])
    prueba('la pregunta del CNAE pierde el 56.30 de la sala y avisa del 47.81',
           '56.30' not in preg and '47.81' in preg and '47.24' in preg, preg[-90:])

    # 4. A6: la bombona sólo en la caseta.
    gas = ev(exc, H_HUM, 'B%d' % G['gas'])
    prueba('en el local con gas: instalación receptora con inspección cada 5 años',
           'Instalación receptora' in gas and '5 años' in gas, gas)
    x4 = ExcelCompiler(ruta)
    ev(x4, H_HUM, 'B%d' % G['gas'])
    ev(x4, H_HUM, 'B%d' % G['prox_gas'])
    x4.set_value("'%s'!B%d" % (H_HUM, G['formato']), FMT_CASETA)
    gas4 = ev(x4, H_HUM, 'B%d' % G['gas'])
    prueba('en la caseta: bombona de menos de 15 kg de utilización móvil y sin inspección',
           'utilización móvil' in gas4 and '15 kg' in gas4
           and ev(x4, H_HUM, 'B%d' % G['prox_gas']) == '', gas4)

    # 5. Alérgenos: freír patatas comparte la freidora y activa la parte A.
    prueba('El Molinete tiene freidora dedicada',
           str(ev(exc, H_ALE, 'B%d' % A['dedicada'])).startswith('Dedicada'))
    x5 = ExcelCompiler(ruta)
    ev(x5, H_ALE, 'B%d' % A['dedicada'])
    ev(x5, H_ALE, 'B%d' % A['ac_a'])
    x5.set_value("'%s'!B%d" % (H_ALE, A['f_ini'] + 2), SI)
    prueba('si fríes patatas: freidora compartida y parte A de la acrilamida',
           str(ev(x5, H_ALE, 'B%d' % A['dedicada'])).startswith('Compartida')
           and 'Te toca la parte A' in ev(x5, H_ALE, 'B%d' % A['ac_a']))

    # 6. Ferias: sin ferias no aplica; con ferias, 11 requisitos pendientes.
    x6 = ExcelCompiler(ruta)
    ev(x6, H_FER, 'B%d' % FE['ver'])
    v6a = ev(x6, H_FER, 'B%d' % FE['ver'])
    x6.set_value("'%s'!B%d" % (H_IAE, I['ferias']), SI)
    v6b = ev(x6, H_FER, 'B%d' % FE['ver'])
    prueba('ferias: la hoja no aplica a El Molinete y se enciende al contestar Sí',
           v6a.startswith('No vas a ferias') and 'pendientes' in v6b, v6b)

    # 7. Checklist: cuenta los trámites y empieza a cero.
    prueba('el checklist cuenta los %d trámites' % len(D.CHECKLIST_LEGAL),
           int(ev(exc, H_CHK, 'D%d' % K_TOTAL)) == len(D.CHECKLIST_LEGAL))
    return ok, fallos


# ==========================================================================
def main():
    D.gate_legal(IDS_LEGALES)
    wb = Workbook()
    wb.remove(wb.active)
    hoja_instrucciones(wb)
    hoja_checklist(wb)
    hoja_iae(wb)
    hoja_regimen(wb)
    hoja_humos(wb)
    hoja_sanitario(wb)
    hoja_ferias(wb)
    hoja_alergenos(wb)
    hoja_prl(wb)
    hoja_formacion(wb)
    hoja_gantt(wb)

    res = cerrar(wb, mapa_celdas())
    print('escrito: %s' % res['ruta'])
    print('hojas: %d · fórmulas: %d · celdas verdes: %d · verdes vacías: %d'
          % (res['hojas'], res['formulas'], res['verdes'], len(res['verdes_vacias'])))
    print('fórmulas que devuelven «sin dato» a propósito: %d' % res['sin_dato'])
    print('notas legales: %d (ids distintos: %d de %d) · etiquetas en el mapa: %d · cruces: 0'
          % (res['notas_legales'], res['ids_con_nota'], len(IDS_LEGALES), res['mapa']))
    print('inject_cache: %s' % res['cache'])
    ok, fallos = demo(res['ruta'])
    for t in ok:
        print('  demo OK    · %s' % t)
    for t in fallos:
        print('  demo FALLA · %s' % t)
    resumen = {'formulas': res['formulas'], 'sin_cache': res['sin_dato'], 'verdes': res['verdes'],
               'verdes_vacias': len(res['verdes_vacias']), 'notas_legales': res['notas_legales'],
               'mapa': res['mapa'], 'demos_ok': not fallos, 'hojas': res['hojas']}
    print('RESUMEN_JSON ' + json.dumps(resumen))
    if fallos:
        raise SystemExit('demos con fallos: %d' % len(fallos))


if __name__ == '__main__':
    main()
