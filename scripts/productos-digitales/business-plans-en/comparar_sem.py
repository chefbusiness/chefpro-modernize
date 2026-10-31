#!/usr/bin/env python3
"""
comparar_sem.py — un xlsx regenerado (dry-run) frente al LIVE de dl/, por CONTENIDO (sesión Claude Code, 31-oct-2026,
resync de los conjuntos EN tras #186). El canon de bytes (`canon.py`) deja de bastar cuando el LIVE se arregló por parche
XML (`parche_xml_xlsx.py`): el regenerado trocea los literales > 255 en otro punto y recalcula el alto de las filas.

    python3 comparar_sem.py <carpeta dry-run con subcarpetas <slug>/> [--dl astro-site/public/dl]

Por fichero: IGUAL-SEM si coinciden pestañas, cada celda (fórmulas tras `fundir_literales`), la caché, number_format,
DV, CF, combinadas, cabeceras/pies y protección; y, de los miembros del zip que difieren en bytes, QUÉ difiere: celdas
cuyo XML cambia (troceo distinto), altos de fila (`ht`) y si queda algo más. rc 1 si algún fichero es DISTINTO-SEM.
"""
import argparse
import glob
import os
import re
import sys
import warnings
import zipfile

warnings.filterwarnings('ignore')
AQUI = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.dirname(AQUI))
from parche_xml_xlsx import fundir_literales                    # noqa: E402
import openpyxl                                                 # noqa: E402

REPO = os.path.abspath(os.path.join(AQUI, '..', '..', '..'))
RX_FECHAS = re.compile(rb'<dcterms:(created|modified)[^>]*>[^<]*</dcterms:\1>')
RX_C = re.compile(r'<c r="([A-Z]+\d+)"[^>]*?(?:/>|>.*?</c>)', re.S)
RX_ROW = re.compile(r'<row r="(\d+)"([^>]*)>')


def F(v):
    return fundir_literales(v) if isinstance(v, str) and v.startswith('=') else v


def dvs(ws):
    return sorted(((str(d.sqref), d.type, F('=' + (d.formula1 or '')), F('=' + (d.formula2 or '')), d.prompt,
                    d.promptTitle, d.error, d.errorTitle, d.allow_blank) for d in ws.data_validations.dataValidation),
                  key=str)


def cfs(ws):
    return sorted(((str(cf.sqref), r.type, r.operator, tuple(F('=' + x) for x in (r.formula or [])), r.text,
                    r.priority, r.dxfId) for cf in ws.conditional_formatting for r in cf.rules), key=str)


def semantica(gen, live):
    difs = []
    a, b = openpyxl.load_workbook(gen), openpyxl.load_workbook(live)
    av, bv = openpyxl.load_workbook(gen, data_only=True), openpyxl.load_workbook(live, data_only=True)
    if a.sheetnames != b.sheetnames:
        return [('pestañas', a.sheetnames, b.sheetnames)]
    for n in a.sheetnames:
        wa, wb, xa, xb = a[n], b[n], av[n], bv[n]
        for k in sorted(set(wa._cells) | set(wb._cells)):
            ca, cb = wa._cells.get(k), wb._cells.get(k)
            va, vb = (F(ca.value) if ca else None), (F(cb.value) if cb else None)
            if va != vb:
                difs.append((n, k, 'valor', va, vb))
            ka, kb = xa._cells.get(k), xb._cells.get(k)
            ka, kb = (ka.value if ka else None), (kb.value if kb else None)
            if ka != kb and not (isinstance(ka, float) and isinstance(kb, float)
                                 and abs(ka - kb) <= 1e-9 * max(1.0, abs(kb))):
                difs.append((n, k, 'caché', ka, kb))
            if ca is not None and cb is not None and ca.number_format != cb.number_format:
                difs.append((n, k, 'formato', ca.number_format, cb.number_format))
        for nom, fa, fb in (('DV', dvs(wa), dvs(wb)), ('CF', cfs(wa), cfs(wb))):
            if fa != fb:
                difs.append((n, nom, [x for x in fa if x not in fb][:2], [x for x in fb if x not in fa][:2]))
        if sorted(map(str, wa.merged_cells.ranges)) != sorted(map(str, wb.merged_cells.ranges)):
            difs.append((n, 'combinadas'))
        for parte in ('oddHeader', 'oddFooter', 'evenHeader', 'evenFooter'):
            for pos in ('left', 'center', 'right'):
                if getattr(getattr(wa, parte), pos).text != getattr(getattr(wb, parte), pos).text:
                    difs.append((n, parte, pos))
        if wa.protection.sheet != wb.protection.sheet:
            difs.append((n, 'protección'))
    return difs


def bytes_distintos(gen, live):
    za, zb = zipfile.ZipFile(gen), zipfile.ZipFile(live)
    hojas = re.findall(r'<sheet [^>]*name="([^"]+)"', zb.read('xl/workbook.xml').decode())
    out = []
    for m in sorted(set(za.namelist()) | set(zb.namelist())):
        if m not in za.namelist() or m not in zb.namelist():
            out.append('%s: falta en uno' % m)
            continue
        x, y = za.read(m), zb.read(m)
        if m == 'docProps/core.xml':
            x, y = RX_FECHAS.sub(b'', x), RX_FECHAS.sub(b'', y)
        if x == y:
            continue
        x, y = x.decode('utf-8', 'replace'), y.decode('utf-8', 'replace')
        s = re.match(r'xl/worksheets/sheet(\d+)\.xml$', m)
        if not s:
            out.append('%s: difiere' % m)
            continue
        ca = {k.group(1): k.group(0) for k in RX_C.finditer(x)}
        cb = {k.group(1): k.group(0) for k in RX_C.finditer(y)}
        ra = {k.group(1): k.group(2) for k in RX_ROW.finditer(x)}
        rb = {k.group(1): k.group(2) for k in RX_ROW.finditer(y)}
        celdas = [k for k in sorted(set(ca) | set(cb)) if ca.get(k) != cb.get(k)]
        filas = ['%s %s→%s' % (k, re.sub(r'\s*customHeight="1"', '', rb.get(k, '')).strip(),
                                re.sub(r'\s*customHeight="1"', '', ra.get(k, '')).strip())
                 for k in sorted(set(ra) | set(rb), key=int) if ra.get(k) != rb.get(k)]
        resto = RX_ROW.sub('<row>', RX_C.sub('', x)) == RX_ROW.sub('<row>', RX_C.sub('', y))
        out.append('%s: XML de celda %s · alto de fila (LIVE→regenerado) %s%s' % (
            hojas[int(s.group(1)) - 1], celdas or '-', filas or '-', '' if resto else ' · Y MÁS'))
    return out


def main():
    ap = argparse.ArgumentParser(description=__doc__.split('\n')[1])
    ap.add_argument('dry')
    ap.add_argument('--dl', default=os.path.join(REPO, 'astro-site', 'public', 'dl'))
    args = ap.parse_args()
    ok = True
    for gen in sorted(glob.glob(os.path.join(args.dry, '**', '*.xlsx'), recursive=True)):
        slug = os.path.basename(os.path.dirname(gen))
        live = os.path.join(args.dl, slug, os.path.basename(gen))
        if not os.path.exists(live):
            print('SIN-LIVE     %s/%s' % (slug, os.path.basename(gen)))
            continue
        difs = semantica(gen, live)
        ok = ok and not difs
        print('%-12s %s/%s · %d difs' % ('IGUAL-SEM' if not difs else 'DISTINTO-SEM', slug, os.path.basename(gen),
                                         len(difs)))
        for d in difs[:20]:
            print('    ', repr(d)[:300])
        for linea in bytes_distintos(gen, live):
            print('      bytes ·', linea[:300])
    sys.exit(0 if ok else 1)


if __name__ == '__main__':
    main()
