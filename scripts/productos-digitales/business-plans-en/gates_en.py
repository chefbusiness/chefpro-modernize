#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
gates_en.py — Gates F2 de Food Truck + Coffee Shop Business Plan Kit (EN), SPEC §8 (G1-G9), sobre la salida de
aplicar_en.py. Copia adaptada de `financial-kit/gates_en.py`. Solo LEE: los 4 xlsx EN, los 4 xlsx ES publicados
(solo lectura), censo_es.json, cifras_caso.json, F2-NOTAS.md y mapas.py. Ningún recuento está escrito a mano: sale
del censo, de los xlsx ES o de mapas.py; las únicas cifras propias son los umbrales D15 (SPEC) y las reglas US de §3.

    python3 gates_en.py                               # contra astro-site/public/dl/<slug>/ y cifras_caso.json
    python3 gates_en.py --dir /tmp/bp-dry             # contra un dry-run (subcarpetas <slug>/ + cifras_caso.json)
    python3 gates_en.py --solo G1,G2 [--mes October] [--json informe.json]
    python3 gates_en.py --autotest [--dir …]          # inyecta defectos en una COPIA y exige que se detecten

Gates:
  G1 paridad (contra el ES): pestañas §2.2; fórmula a fórmula = ES tras E2 + pestañas + CLAVES salvo FORMULAS_EN
     (celda a celda); tipo de cada celda (fórmula / número / texto / vacía) igual salvo E1 y los VALORES_EN vacíos;
     DV (sqref, tipo, operador, flags, fórmulas por CLAVES/pestañas, mensajes presentes); CF (sqref, tipo,
     prioridad, formato, fórmula y `text` por CLAVES); merges; paneles; filas/columnas ocultas; protección (sin
     contraseña, mismos flags); verdes = desbloqueadas = las del ES; tareas por fase (68 / 75);
  G2 restos: cero español (mapas.restos_espanol), «€», «m²», « », no latinos, coma decimal, normativa ES, «v1.1»,
     «Fichero», IDs internos en celdas, literales, DV, CF, pies y docProps; formatos sin «€»; DV con título ≤ 32 y
     mensaje ≤ 255;
  G3 formato: Letter en las 32 hojas, pie EN donde el ES tenía pie, línea de versión D20, aviso D21 (E1) con el
     estilo de la versión, docProps D20, formatos D8 celda a celda, marca «More templates…»;
  G4 claves: ítems de lista ∈ CLAVES; cada valor bajo una DV de lista = un ítem; cada número bajo una DV numérica
     dentro de sus límites; literales de fórmula con letras ∈ CLAVES (o celda de FORMULAS_EN);
  G5 CF: tokens EN inyectivos por rango, ∈ CLAVES, ningún literal cambia de color; mapas.cruzar_censo sin errores;
  G6 cálculo: caché en toda fórmula, cero errores; caché EN = caché de REFERENCIA (xlsx ES con E2 + VALORES_EN +
     CALIBRACION_D15 recalculado con pycel) celda a celda; E2 → input tax de compras = 0 y sales tax remitido cada
     trimestre = cobrado; nóminas constantes con 12 pagas; D15 (DSCR mín ≥ 1.25, saldo mín > 0, cobertura ≥ 98 %,
     sueldos ≥ suelo, holgura ≥ 15 %); préstamo = el que propone Financing; cifras_caso.json = caché;
  G7 reglas: VALORES_EN / CALIBRACION_D15 escritos; cada calibración citada en F2-NOTAS.md; reglas US de §3 sin
     calibrar; ningún texto con «SBA minimum» ni «approval guaranteed»;
  G8 docx (si existen): ensamblar_docx.g8 con cifras_caso.json;
  G9 censo-entregables.py --only <slug> --fail y gate-no-latinos.py --only <slug> en 0.
Sale con código 1 si algún gate falla.
"""
import argparse
import copy
import json
import logging
import os
import re
import shutil
import subprocess
import sys
import tempfile
import time
import zipfile
from collections import OrderedDict

AQUI = os.path.dirname(os.path.abspath(__file__))
SCRIPTS = os.path.dirname(AQUI)
REPO = os.path.dirname(os.path.dirname(SCRIPTS))
sys.path.insert(0, AQUI)
sys.path.insert(1, SCRIPTS)

import openpyxl                                              # noqa: E402

import mapas                                                 # noqa: E402
import aplicar_en as A                                       # noqa: E402

logging.disable(logging.CRITICAL)

CENSO_SCRIPT = os.path.join(SCRIPTS, 'censo-entregables.py')
NO_LATINOS = os.path.join(SCRIPTS, 'gate-no-latinos.py')
TOL = 0.006
ERRORES = ('#REF!', '#VALUE!', '#NAME?', '#DIV/0!', '#N/A', '#NUM!', '#NULL!')
CLAVES_EN = set(mapas.CLAVES.values())
LIMITE_DV = {'promptTitle': 32, 'errorTitle': 32, 'prompt': 255, 'error': 255}
RX_RESTOS_EXTRA = re.compile(r'v1\.1|\b[Ff]ichero\b|Versión \d|\bprovincial\b')
PROHIBIDO = re.compile(r'SBA minimum|minimum required by the SBA|approval (?:is )?guaranteed|guaranteed approval',
                       re.I)
# reglas US de SPEC §3 (research §5/§7): valor exigido; nunca se calibran
REGLAS_CELDA = OrderedDict([('B20', 0.10), ('B21', 12), ('B22', 15080), ('B32', 0.10), ('B34', 0), ('B37', 0.25),
                            ('B38', 0.25), ('B39', 0.08), ('B40', 0.08), ('B41', 0), ('B63', 0.08), ('B65', 1.15),
                            ('B66', 1.25)])


# ==========================================================================
# Utilidades
# ==========================================================================
class Resultado:
    def __init__(self, silencioso=False):
        self.fallos = OrderedDict()
        self.info = OrderedDict()
        self.silencioso = silencioso

    def f(self, gate, msg):
        self.fallos.setdefault(gate, []).append(msg)

    def i(self, gate, msg):
        self.info.setdefault(gate, []).append(msg)

    def gates_rojos(self):
        return [g for g, v in self.fallos.items() if v]


_CACHE = {}


def cargar(path, data_only=False):
    k = (path, data_only, os.path.getmtime(path))
    if k not in _CACHE:
        _CACHE[k] = openpyxl.load_workbook(path, data_only=data_only)
    return _CACHE[k]


def carpetas_de(base):
    if base:
        return OrderedDict((p, os.path.join(base, mapas.PLANES[p]['slug'])) for p in mapas.PLANES)
    return OrderedDict((p, A.destino(p)) for p in mapas.PLANES)


def ruta_en(carpetas, corto):
    return os.path.join(carpetas[mapas.PLAN_DE[corto]], mapas.LIBROS[corto][2])


def ruta_es(corto):
    return mapas.ruta_es(corto, REPO)


def es_num(v):
    return isinstance(v, (int, float)) and not isinstance(v, bool)


def es_formula(c):
    return c.data_type == 'f' or (isinstance(c.value, str) and c.value.startswith('='))


def lits(f):
    return [m.group(0)[1:-1].replace('""', '"') for m in A.RX_STR.finditer(f or '')]


def celdas(sqref):
    out = []
    for rg in str(sqref).split():
        out.extend(A.rango(rg))
    return out


def es_verde(c):
    return c.fill is not None and c.fill.fill_type == 'solid' and str(c.fill.fgColor.rgb)[-6:] == A.VERDE


def proteccion(ws):
    p = ws.protection
    flags = ('sheet', 'formatCells', 'formatColumns', 'formatRows', 'insertColumns', 'insertRows', 'insertHyperlinks',
             'deleteColumns', 'deleteRows', 'selectLockedCells', 'selectUnlockedCells', 'sort', 'autoFilter',
             'pivotTables', 'objects', 'scenarios')
    return tuple((f, bool(getattr(p, f, None))) for f in flags) + (('password', p.password or None),)


def dxf_firma(regla):
    d = regla.dxf
    if d is None:
        return None
    fill = None
    if d.fill is not None:
        fill = (str(getattr(d.fill.fgColor, 'rgb', None)), str(getattr(d.fill.bgColor, 'rgb', None)))
    font = None
    if d.font is not None and d.font.color is not None:
        font = (str(d.font.color.rgb), bool(d.font.b))
    return (fill, font)


def textos_de_formula_en(f):
    return [x for x in lits(f) if any(ch.isalpha() for ch in x)]


def valores_esperados(corto, cifras):
    """(hoja ES, celda) → valor exigido en el EN (VALORES_EN + CALIBRACION_D15 + préstamo de cifras_caso)."""
    plan = mapas.PLAN_DE[corto]
    prest = {plan: int(round(cifras[plan]['prestamo']['valor']))} if cifras and plan in cifras else {plan: None}
    return A.valores_libro(corto, prest)


def cargar_cifras(base):
    p = os.path.join(base, 'cifras_caso.json') if base else os.path.join(AQUI, 'cifras_caso.json')
    if not os.path.exists(p):
        return None, p
    return json.load(open(p, encoding='utf-8')), p


# ==========================================================================
# G1 · paridad
# ==========================================================================
def gateG1(R, carpetas, datos):
    G = 'G1'
    nuevas = {(c, h, x) for (c, h, x) in mapas.TEXTOS_NUEVOS}
    for corto in mapas.LIBROS:
        wes, wen = cargar(ruta_es(corto)), cargar(ruta_en(carpetas, corto))
        esperadas = [mapas.hoja_en(corto, h) for h in wes.sheetnames]
        if wen.sheetnames != esperadas:
            R.f(G, '%s: pestañas %r ≠ %r' % (corto, wen.sheetnames, esperadas))
            continue
        vac = {(h, x) for (h, x), v in A.valores_libro(corto, {mapas.PLAN_DE[corto]: 0}).items() if v is None} \
            if mapas.TIPO_DE[corto] == 'plan' else set()
        n_f = 0
        for h_es in wes.sheetnames:
            a, b = wes[h_es], wen[mapas.hoja_en(corto, h_es)]
            coords = {c.coordinate for c in a._cells.values() if c.value is not None} | \
                     {c.coordinate for c in b._cells.values() if c.value is not None}
            for x in sorted(coords):
                ca, cb = a[x], b[x]
                va, vb = ca.value, cb.value
                if va is None:
                    if (corto, h_es, x) not in nuevas:
                        R.f(G, '%s %s!%s: celda nueva %r (solo E1 puede añadir)' % (corto, h_es, x, str(vb)[:40]))
                    continue
                if es_formula(ca):
                    n_f += 1
                    try:
                        esp = A.transformar_es(corto, h_es, x, va)
                    except Exception as e:                       # noqa: BLE001
                        R.f(G, '%s %s!%s: no se puede transformar la fórmula ES: %s' % (corto, h_es, x, e))
                        continue
                    if vb != esp:
                        R.f(G, '%s %s!%s: fórmula %r ≠ esperada %r' % (corto, h_es, x, str(vb)[:70], esp[:70]))
                elif es_num(va):
                    if not (es_num(vb) or (vb is None and (h_es, x) in vac)):
                        R.f(G, '%s %s!%s: número en el ES y %r en el EN' % (corto, h_es, x, vb))
                elif isinstance(va, str):
                    if not isinstance(vb, str) or es_formula(cb) or not vb.strip():
                        R.f(G, '%s %s!%s: texto en el ES y %r en el EN' % (corto, h_es, x, str(vb)[:40]))
            if sorted(map(str, a.merged_cells.ranges)) != sorted(map(str, b.merged_cells.ranges)):
                R.f(G, '%s %s: merges distintos' % (corto, h_es))
            if a.freeze_panes != b.freeze_panes:
                R.f(G, '%s %s: paneles %s ≠ %s' % (corto, h_es, b.freeze_panes, a.freeze_panes))
            if proteccion(a) != proteccion(b):
                R.f(G, '%s %s: protección distinta %s' % (corto, h_es, [x for x, y in zip(proteccion(b), proteccion(a))
                                                                       if x != y]))
            oc = lambda ws: (sorted(k for k, d in ws.column_dimensions.items() if d.hidden),       # noqa: E731
                             sorted(k for k, d in ws.row_dimensions.items() if d.hidden))
            if oc(a) != oc(b):
                R.f(G, '%s %s: filas/columnas ocultas distintas' % (corto, h_es))
            verdes_a = {c.coordinate for c in a._cells.values() if es_verde(c)}
            verdes_b = {c.coordinate for c in b._cells.values() if es_verde(c)}
            desb_a = {c.coordinate for c in a._cells.values() if c.protection is not None and not c.protection.locked}
            desb_b = {c.coordinate for c in b._cells.values() if c.protection is not None and not c.protection.locked}
            if verdes_a != verdes_b:
                R.f(G, '%s %s: verdes distintas %s' % (corto, h_es, sorted(verdes_a ^ verdes_b)[:5]))
            if desb_a != desb_b:
                R.f(G, '%s %s: desbloqueadas distintas %s' % (corto, h_es, sorted(desb_a ^ desb_b)[:5]))
            # DV
            T = A.Transformador(mapas.HOJAS_POR_LIBRO[corto])
            dva = {str(d.sqref): d for d in a.data_validations.dataValidation}
            dvb = {str(d.sqref): d for d in b.data_validations.dataValidation}
            if set(dva) != set(dvb):
                R.f(G, '%s %s: DV %s ≠ %s' % (corto, h_es, sorted(dvb), sorted(dva)))
            for sq in set(dva) & set(dvb):
                da, db = dva[sq], dvb[sq]
                for at in ('type', 'operator', 'allow_blank', 'showErrorMessage', 'showInputMessage', 'errorStyle',
                           'showDropDown'):
                    if getattr(da, at) != getattr(db, at):
                        R.f(G, '%s %s DV %s: %s %r ≠ %r' % (corto, h_es, sq, at, getattr(db, at), getattr(da, at)))
                f1 = da.formula1 or ''
                if da.type == 'list' and f1.startswith('"'):
                    esp = '"' + ','.join(mapas.CLAVES.get(x, x) for x in f1.strip('"').split(',')) + '"'
                else:
                    esp = T.formula(f1) if f1 else f1
                if (db.formula1 or '') != esp:
                    R.f(G, '%s %s DV %s: formula1 %r ≠ %r' % (corto, h_es, sq, db.formula1, esp))
                if (da.formula2 or None) and (db.formula2 != T.formula(da.formula2)):
                    R.f(G, '%s %s DV %s: formula2 distinta' % (corto, h_es, sq))
                for at in ('error', 'errorTitle', 'prompt', 'promptTitle'):
                    if bool(getattr(da, at)) != bool(getattr(db, at)):
                        R.f(G, '%s %s DV %s: %s presente en uno solo' % (corto, h_es, sq, at))
            # CF
            def firma_cf(ws, trans):
                out = []
                for cf in ws.conditional_formatting:
                    for r in cf.rules:
                        out.append((str(cf.sqref), r.type, r.priority, r.operator, r.stopIfTrue, dxf_firma(r),
                                    tuple(trans(x) for x in (r.formula or [])),
                                    mapas.CLAVES.get(r.text, r.text) if trans is not None and r.text else r.text))
                return sorted(out, key=repr)
            fa = firma_cf(a, T.formula)
            fb = firma_cf(b, lambda x: x)
            fb = [t[:7] + (t[7],) for t in fb]
            if fa != fb:
                dif = [x for x in fb if x not in fa][:2]
                R.f(G, '%s %s: CF distinta (%d ≠ %d reglas) %r' % (corto, h_es, len(fb), len(fa), dif))
        n_cen = sum(info['totales'].get('celdas_formula', 0) for info in [A.censo_libro(datos, corto)])
        if n_f != n_cen:
            R.f(G, '%s: %d fórmulas ≠ censo %d' % (corto, n_f, n_cen))
        if mapas.TIPO_DE[corto] == 'check':
            fases = []
            for h_es in wes.sheetnames[:-1]:
                ws = wen[mapas.hoja_en(corto, h_es)]
                fases.append(sum(1 for c in ws['A'] if c.value in ('☐', '✓', 'N/A')))
            if fases != mapas.TAREAS_POR_FASE[corto]:
                R.f(G, '%s: tareas por fase %s ≠ %s' % (corto, fases, mapas.TAREAS_POR_FASE[corto]))
            else:
                R.i(G, '%s: %d tareas %s' % (corto, sum(fases), fases))
        R.i(G, '%s: %d hojas · %d fórmulas' % (corto, len(wen.sheetnames), n_f))


# ==========================================================================
# G2 · restos
# ==========================================================================
def piezas_de_texto(corto, wb):
    """(lugar, texto) de todo el texto visible del libro EN."""
    out = []
    for ws in wb.worksheets:
        for c in ws._cells.values():
            v = c.value
            if not isinstance(v, str):
                continue
            if es_formula(c):
                for x in textos_de_formula_en(v):
                    out.append(('%s %s!%s (literal)' % (corto, ws.title, c.coordinate), x))
            else:
                out.append(('%s %s!%s' % (corto, ws.title, c.coordinate), v))
        for dv in ws.data_validations.dataValidation:
            for at in ('error', 'errorTitle', 'prompt', 'promptTitle'):
                v = getattr(dv, at)
                if v:
                    out.append(('%s %s DV %s %s' % (corto, ws.title, dv.sqref, at), v))
            if dv.type == 'list' and (dv.formula1 or '').startswith('"'):
                for it in dv.formula1.strip('"').split(','):
                    out.append(('%s %s DV %s ítem' % (corto, ws.title, dv.sqref), it))
        for cf in ws.conditional_formatting:
            for r in cf.rules:
                if r.text:
                    out.append(('%s %s CF %s' % (corto, ws.title, cf.sqref), r.text))
        for parte in ('oddHeader', 'oddFooter', 'evenHeader', 'evenFooter', 'firstHeader', 'firstFooter'):
            hf = getattr(ws, parte)
            for pos in ('left', 'center', 'right'):
                t = getattr(hf, pos).text
                if t:
                    out.append(('%s %s %s.%s' % (corto, ws.title, parte, pos), t))
    p = wb.properties
    for k in ('title', 'subject', 'keywords', 'description', 'category', 'creator', 'lastModifiedBy'):
        v = getattr(p, k)
        if v:
            out.append(('%s docProps %s' % (corto, k), v))
    return out


def gateG2(R, carpetas, datos):
    G = 'G2'
    n = 0
    for corto in mapas.LIBROS:
        wb = cargar(ruta_en(carpetas, corto))
        for lugar, t in piezas_de_texto(corto, wb):
            n += 1
            for motivo, ctx in mapas.restos_espanol(t):
                R.f(G, '%s: %s · «%s»' % (lugar, motivo, ctx.strip()[:70]))
            if mapas.no_latinos(t):
                R.f(G, '%s: no latinos %s' % (lugar, mapas.no_latinos(t)))
            m = RX_RESTOS_EXTRA.search(t)
            if m:
                R.f(G, '%s: resto del ES «%s»' % (lugar, m.group(0)))
            if mapas.RX_ID_INTERNO.search(t):
                R.f(G, '%s: ID interno «%s»' % (lugar, mapas.RX_ID_INTERNO.search(t).group(0)))
        for ws in wb.worksheets:
            for c in ws._cells.values():
                if '€' in (c.number_format or ''):
                    R.f(G, '%s %s!%s: formato con €: %r' % (corto, ws.title, c.coordinate, c.number_format))
            for dv in ws.data_validations.dataValidation:
                for at, lim in LIMITE_DV.items():
                    v = getattr(dv, at)
                    if v and len(v) > lim:
                        R.f(G, '%s %s DV %s: %s de %d caracteres > %d' % (corto, ws.title, dv.sqref, at, len(v), lim))
        for f in wb._number_formats:
            if '€' in f:
                R.f(G, '%s: formato propio con € en el registro: %r' % (corto, f))
    R.i(G, '%d piezas de texto revisadas' % n)


# ==========================================================================
# G3 · formato
# ==========================================================================
def gateG3(R, carpetas, datos, mes):
    G = 'G3'
    n_hojas = 0
    for corto in mapas.LIBROS:
        plan = mapas.PLAN_DE[corto]
        wes, wen = cargar(ruta_es(corto)), cargar(ruta_en(carpetas, corto))
        for h_es in wes.sheetnames:
            a, b = wes[h_es], wen[mapas.hoja_en(corto, h_es)]
            n_hojas += 1
            if b.page_setup.paperSize not in (1, '1'):
                R.f(G, '%s %s: paperSize %r (Letter = 1)' % (corto, b.title, b.page_setup.paperSize))
            for parte in ('oddHeader', 'oddFooter', 'evenHeader', 'evenFooter', 'firstHeader', 'firstFooter'):
                for pos in ('left', 'center', 'right'):
                    ta, tb = getattr(getattr(a, parte), pos).text, getattr(getattr(b, parte), pos).text
                    if ta and tb != mapas.FOOTER_EN:
                        R.f(G, '%s %s %s.%s: %r ≠ pie EN' % (corto, b.title, parte, pos, tb))
                    if not ta and tb:
                        R.f(G, '%s %s %s.%s nuevo: %r' % (corto, b.title, parte, pos, tb))
            for c in a._cells.values():
                fa, fb = c.number_format, b[c.coordinate].number_format
                esp = mapas.FORMATOS_EN.get(fa, fa)
                if fb != esp:
                    R.f(G, '%s %s!%s: formato %r ≠ %r' % (corto, b.title, c.coordinate, fb, esp))
        ins = wen['Instructions']
        fila = mapas.FILA_VERSION[corto]
        ver = mapas.version_en(plan, mes)
        if ins['A%d' % fila].value != ver:
            R.f(G, '%s: Instructions!A%d = %r ≠ %r' % (corto, fila, ins['A%d' % fila].value, ver))
        for (c, h, x), (texto, estilo) in mapas.TEXTOS_NUEVOS.items():
            if c != corto:
                continue
            if ins[x].value != texto:
                R.f(G, '%s: aviso D21 en Instructions!%s = %r' % (corto, x, ins[x].value))
            sa, sb = list(ins[x]._style), list(ins[estilo]._style)
            if sa[:5] + sa[6:] != sb[:5] + sb[6:] or not ins[x].alignment.wrap_text:
                R.f(G, '%s: el aviso D21 no lleva el estilo de la línea de versión (con ajuste de texto)' % corto)
            if (ins.row_dimensions[ins[x].row].height or 0) < 2 * 11:
                R.f(G, '%s: la fila del aviso D21 no tiene alto para su texto ajustado' % corto)
        marca = mapas.FIJOS['Más plantillas y kits del catálogo en aichef.pro/productos-digitales']
        if not any(c.value == marca for c in ins['A']):
            R.f(G, '%s: falta la línea de marca %r' % (corto, marca))
        p = wen.properties
        for k, v in mapas.docprops_en(corto).items():
            if getattr(p, k) != v:
                R.f(G, '%s docProps %s = %r ≠ %r' % (corto, k, getattr(p, k), v))
    R.i(G, '%d hojas en Letter con pie, versión y aviso' % n_hojas)


# ==========================================================================
# G4 · claves y límites de DV
# ==========================================================================
def _cumple(v, op, a, b):
    if op in (None, 'between'):
        return a <= v <= b if b is not None else v >= a
    return {'greaterThanOrEqual': v >= a, 'greaterThan': v > a, 'lessThanOrEqual': v <= a, 'lessThan': v < a,
            'equal': v == a, 'notEqual': v != a, 'notBetween': not (a <= v <= (b if b is not None else a))}[op]


def _num_lit(x):
    try:
        return float(x)
    except (TypeError, ValueError):
        return None


def gateG4(R, carpetas, datos):
    G = 'G4'
    n_val = 0
    for corto in mapas.LIBROS:
        wen = cargar(ruta_en(carpetas, corto))
        for ws in wen.worksheets:
            for dv in ws.data_validations.dataValidation:
                f1 = dv.formula1 or ''
                if dv.type == 'list' and f1.startswith('"'):
                    items = f1.strip('"').split(',')
                    for it in items:
                        if any(ch.isalpha() for ch in it) and it not in CLAVES_EN:
                            R.f(G, '%s %s DV %s: ítem %r fuera de CLAVES' % (corto, ws.title, dv.sqref, it))
                    for x in celdas(dv.sqref):
                        v = ws[x].value
                        if v is None or es_formula(ws[x]):
                            continue
                        n_val += 1
                        if str(v) not in items:
                            R.f(G, '%s %s!%s: %r no es un ítem de su lista %r' % (corto, ws.title, x, v, items))
                elif dv.type in ('decimal', 'whole'):
                    a_, b_ = _num_lit(dv.formula1), _num_lit(dv.formula2)
                    if a_ is None:
                        continue
                    for x in celdas(dv.sqref):
                        v = ws[x].value
                        if not es_num(v):
                            continue
                        n_val += 1
                        if not _cumple(v, dv.operator, a_, b_):
                            R.f(G, '%s %s!%s: %r fuera de su DV (%s %s %s)' % (corto, ws.title, x, v, dv.operator,
                                                                             dv.formula1, dv.formula2))
            h_inv = {v: k for k, v in mapas.HOJAS_POR_LIBRO[corto].items()}
            for c in ws._cells.values():
                if not es_formula(c):
                    continue
                if (corto, h_inv[ws.title], c.coordinate) in A._CUBIERTAS:
                    continue
                for x in textos_de_formula_en(c.value):
                    if x not in CLAVES_EN:
                        R.f(G, '%s %s!%s: literal %r fuera de CLAVES' % (corto, ws.title, c.coordinate, x))
    R.i(G, '%d valores bajo DV comprobados' % n_val)


# ==========================================================================
# G5 · CF
# ==========================================================================
def gateG5(R, carpetas, datos):
    G = 'G5'
    for e in mapas.cruzar_censo(os.path.join(AQUI, 'censo_es.json')):
        R.f(G, 'cruzar_censo: ' + e)
    n = 0
    for corto in mapas.LIBROS:
        wes, wen = cargar(ruta_es(corto)), cargar(ruta_en(carpetas, corto))
        for h_es in wes.sheetnames:
            a, b = wes[h_es], wen[mapas.hoja_en(corto, h_es)]
            toks_a, toks_b = OrderedDict(), OrderedDict()
            for ws, d in ((a, toks_a), (b, toks_b)):
                for cf in ws.conditional_formatting:
                    for r in cf.rules:
                        if r.text:
                            d.setdefault(str(cf.sqref), []).append(r.text)
            for sq, tb in toks_b.items():
                n += 1
                ta = toks_a.get(sq, [])
                if len({t.lower() for t in tb}) != len(tb):
                    R.f(G, '%s %s CF %s: tokens EN no inyectivos %r' % (corto, b.title, sq, tb))
                for t in tb:
                    if any(ch.isalpha() for ch in t) and t not in CLAVES_EN:
                        R.f(G, '%s %s CF %s: token %r fuera de CLAVES' % (corto, b.title, sq, t))
                # ningún literal cambia de color: mismas reglas casadas antes y después
                for x in celdas(sq):
                    fa, fb = a[x].value, b[x].value
                    if not (isinstance(fa, str) and isinstance(fb, str)):
                        continue
                    la, lb = lits(fa), lits(fb)
                    if len(la) != len(lb):
                        continue
                    for xa, xb in zip(la, lb):
                        ma = [i for i, t in enumerate(ta) if t.lower() in xa.lower()]
                        mb = [i for i, t in enumerate(tb) if t.lower() in xb.lower()]
                        if ma != mb:
                            R.f(G, '%s %s!%s: el literal %r cambia de color (%s → %s)' % (corto, b.title, x, xb, ma, mb))
    R.i(G, '%d rangos con tokens de CF' % n)


# ==========================================================================
# G6 · cálculo
# ==========================================================================
_REF = {}


def referencia(corto, cifras):
    """xlsx ES con E2 + VALORES_EN + CALIBRACION_D15 (en espacio ES), recalculado con pycel → data_only."""
    if corto in _REF:
        return _REF[corto]
    tmp = tempfile.mkdtemp(prefix='bp-ref-')
    p = os.path.join(tmp, 'ref-%s.xlsx' % corto)
    wb = openpyxl.load_workbook(ruta_es(corto))
    for ws in wb.worksheets:
        for c in ws._cells.values():
            if es_formula(c):
                c.value = mapas.parchear_formula_es(corto, ws.title, c.coordinate, c.value)
    for (h, x), v in valores_esperados(corto, cifras).items():
        wb[h][x].value = v
    wb.save(p)
    A.inyectar(p)
    _REF[corto] = openpyxl.load_workbook(p, data_only=True)
    shutil.rmtree(tmp, ignore_errors=True)
    return _REF[corto]


def gateG6(R, carpetas, datos, cifras, cifras_path):
    G = 'G6'
    for corto in mapas.LIBROS:
        p = ruta_en(carpetas, corto)
        wf, wv = cargar(p), cargar(p, data_only=True)
        sin = err = 0
        for ws in wf.worksheets:
            for c in ws._cells.values():
                if not es_formula(c):
                    continue
                v = wv[ws.title][c.coordinate].value
                if v is None and '""' not in c.value:
                    sin += 1
                    if sin <= 3:
                        R.f(G, '%s %s!%s: fórmula sin caché' % (corto, ws.title, c.coordinate))
                if isinstance(v, str) and v in ERRORES:
                    err += 1
                    R.f(G, '%s %s!%s: error %s en la caché' % (corto, ws.title, c.coordinate, v))
        if sin > 3:
            R.f(G, '%s: %d fórmulas sin caché en total' % (corto, sin))
        # caché EN = referencia ES recalculada
        if mapas.TIPO_DE[corto] == 'plan' and cifras:
            ref = referencia(corto, cifras)
            dif = 0
            for h_es in ref.sheetnames:
                wr, we = ref[h_es], wv[mapas.hoja_en(corto, h_es)]
                for c in wf[mapas.hoja_en(corto, h_es)]._cells.values():
                    if not es_formula(c):
                        continue
                    vr, ve = wr[c.coordinate].value, we[c.coordinate].value
                    if es_num(vr) or es_num(ve):
                        if not (es_num(vr) and es_num(ve)) or abs(vr - ve) > TOL * max(1.0, abs(vr)) * 0.01 + 1e-6:
                            dif += 1
                            if dif <= 5:
                                R.f(G, '%s %s!%s: caché EN %r ≠ referencia %r' % (corto, h_es, c.coordinate, ve, vr))
                    elif isinstance(vr, str) and vr in mapas.CLAVES:
                        if ve != mapas.CLAVES[vr]:
                            dif += 1
                            R.f(G, '%s %s!%s: %r ≠ CLAVES[%r]' % (corto, h_es, c.coordinate, ve, vr))
                    elif bool(vr) != bool(ve):
                        dif += 1
                        R.f(G, '%s %s!%s: texto calculado vacío en uno solo (%r / %r)' % (corto, h_es, c.coordinate,
                                                                                       str(ve)[:30], str(vr)[:30]))
            if dif > 5:
                R.f(G, '%s: %d celdas de caché distintas de la referencia' % (corto, dif))
            R.i(G, '%s: caché = referencia ES recalculada (%d diferencias)' % (corto, dif))
        if mapas.TIPO_DE[corto] != 'plan':
            continue
        plan = mapas.PLAN_DE[corto]
        leer = lambda h, x, _c=corto, _wv=wv: _wv[mapas.hoja_en(_c, h)][x].value       # noqa: E731
        cen = A.censo_libro(datos, corto)
        # E2: input tax de las compras = 0 (G13/G14 apuntan a B41 = 0)
        for x in ('G13', 'G14'):
            if leer('PyG 3 Años', x) != 0:
                R.f(G, '%s 3-Year P&L!%s = %r (E2: input tax de compras = 0)' % (corto, x, leer('PyG 3 Años', x)))
        # sales tax: memoria de input tax = 0; remitido cada trimestre = cobrado del trimestre
        f16 = A.fila_de(cen, 'Tesorería 12 meses', 'IVA repercutido del mes')
        f17 = A.fila_de(cen, 'Tesorería 12 meses', 'IVA soportado del mes')
        f20 = A.fila_de(cen, 'Tesorería 12 meses', 'Pago del IVA')
        f13 = A.fila_de(cen, 'Tesorería 12 meses', 'Nóminas')
        cols = 'BCDEFGHIJKLM'
        for col in cols:
            v = leer('Tesorería 12 meses', '%s%d' % (col, f17))
            if es_num(v) and abs(v) > 0.005:
                R.f(G, '%s cash flow %s%d: input tax %r ≠ 0' % (corto, col, f17, v))
        for pago, meses in (('E', 'BCD'), ('H', 'EFG'), ('K', 'HIJ')):
            cobrado = sum(leer('Tesorería 12 meses', '%s%d' % (m, f16)) or 0 for m in meses)
            v = leer('Tesorería 12 meses', '%s%d' % (pago, f20))
            if not es_num(v) or abs(-v - cobrado) > 0.01:
                R.f(G, '%s cash flow %s%d: remitido %r ≠ cobrado del trimestre %.2f' % (corto, pago, f20, v, cobrado))
        nom = [leer('Tesorería 12 meses', '%s%d' % (col, f13)) for col in cols]
        if len({round(x, 4) for x in nom if es_num(x)}) != 1:
            R.f(G, '%s cash flow fila %d: nóminas no constantes con 12 pagas %r' % (corto, f13, nom[:3]))
        # D15
        m, fallos = A.d15(leer, plan, datos)
        for x in fallos:
            R.f(G, '%s D15: %s' % (corto, x))
        R.i(G, '%s D15: DSCR mín %.2f · saldo mín %.0f · cobertura %.1f%% · holgura %.1f%%' % (
            corto, m['dscr_min'], m['saldo_minimo'], m['cobertura_horas'] * 100, m['holgura_caja_pct'] * 100))
        # préstamo
        prop = A.prestamo_propuesto(leer, cen)
        if leer('0. Supuestos', 'B31') != prop:
            R.f(G, '%s: préstamo B31 %r ≠ el que propone Financing %r' % (corto, leer('0. Supuestos', 'B31'), prop))
        # cifras_caso.json = caché
        if not cifras or plan not in cifras:
            R.f(G, 'falta %s o su plan %s' % (cifras_path, plan))
            continue
        nuevo = mapas.calcular_cifras(plan, leer, datos['censo']['tokens_docx'][plan])
        for t, d in nuevo.items():
            c = cifras[plan].get(t)
            if c is None:
                R.f(G, 'cifras_caso.json %s: falta %s' % (plan, t))
                continue
            if c.get('texto') != d['texto']:
                R.f(G, 'cifras_caso.json %s.%s: %r ≠ caché %r' % (plan, t, c.get('texto'), d['texto']))
        sobra = set(cifras[plan]) - set(nuevo)
        if sobra:
            R.f(G, 'cifras_caso.json %s: tokens que ya no existen %r' % (plan, sorted(sobra)))


# ==========================================================================
# G7 · reglas
# ==========================================================================
def gateG7(R, carpetas, datos, cifras):
    G = 'G7'
    notas = open(os.path.join(AQUI, 'F2-NOTAS.md'), encoding='utf-8').read()
    for (c, h, x), (v, motivo) in mapas.CALIBRACION_D15.items():
        cita = '%s %s!%s' % (c, h, x)
        if cita not in notas:
            R.f(G, 'CALIBRACION_D15 %s sin explicar en F2-NOTAS.md' % cita)
        if h == '0. Supuestos' and x in REGLAS_CELDA:
            R.f(G, 'CALIBRACION_D15 toca una regla US de §3: %s' % cita)
    n = 0
    for corto in (k for k in mapas.LIBROS if mapas.TIPO_DE[k] == 'plan'):
        wen = cargar(ruta_en(carpetas, corto))
        for (h, x), v in valores_esperados(corto, cifras).items():
            n += 1
            real = wen[mapas.hoja_en(corto, h)][x].value
            if v is None:
                if real is not None:
                    R.f(G, '%s %s!%s: %r y se esperaba vacía' % (corto, h, x, real))
            elif not es_num(real) or abs(real - v) > 1e-9:
                R.f(G, '%s %s!%s: %r ≠ %r (VALORES_EN / CALIBRACION_D15)' % (corto, h, x, real, v))
        for x, v in REGLAS_CELDA.items():
            real = wen['0. Assumptions'][x].value
            if not es_num(real) or abs(real - v) > 1e-9:
                R.f(G, '%s 0. Assumptions!%s = %r ≠ regla US %r' % (corto, x, real, v))
    for corto in mapas.LIBROS:
        for lugar, t in piezas_de_texto(corto, cargar(ruta_en(carpetas, corto))):
            if PROHIBIDO.search(t):
                R.f(G, '%s: afirmación prohibida «%s»' % (lugar, PROHIBIDO.search(t).group(0)))
    R.i(G, '%d valores del caso comprobados · %d calibraciones citadas' % (n, len(mapas.CALIBRACION_D15)))


# ==========================================================================
# G8 · docx · G9 · gates del repo
# ==========================================================================
def gateG8(R, carpetas, datos, cifras):
    G = 'G8'
    import ensamblar_docx as E
    for plan in mapas.PLANES:
        p = os.path.join(carpetas[plan], mapas.DOCX[plan][1])
        if not os.path.exists(p):
            R.i(G, '%s: sin docx todavía (%s)' % (plan, os.path.basename(p)))
            continue
        for x in E.g8(plan, p, (cifras or {}).get(plan)):
            R.f(G, '%s: %s' % (plan, x))
        R.i(G, '%s: docx validado' % plan)


def gateG9(R, carpetas, datos):
    G = 'G9'
    for plan, carpeta in carpetas.items():
        slug = mapas.PLANES[plan]['slug']
        real = os.path.abspath(carpeta) == os.path.abspath(A.destino(plan))
        only = slug if real else carpeta
        extra = [] if real else ['--letter']
        for cmd in ([sys.executable, CENSO_SCRIPT, '--only', only, '--fail', '--quiet'] + extra,
                    [sys.executable, NO_LATINOS, '--only', only]):
            r = subprocess.run(cmd, stdout=subprocess.PIPE, stderr=subprocess.STDOUT, universal_newlines=True,
                               cwd=REPO)
            ult = (r.stdout.strip().splitlines() or [''])[-1]
            if r.returncode != 0:
                R.f(G, '%s %s: código %d · %s' % (slug, os.path.basename(cmd[1]), r.returncode, r.stdout[-400:]))
            else:
                R.i(G, '%s %s: %s' % (slug, os.path.basename(cmd[1]), ult[:120]))


# ==========================================================================
# Ejecución
# ==========================================================================
def ejecutar(base, mes, solo=None, silencioso=False, g9=True):
    datos = A.cargar_datos()
    carpetas = carpetas_de(base)
    cifras, cifras_path = cargar_cifras(base)
    R = Resultado(silencioso)
    gates = OrderedDict([
        ('G1', lambda: gateG1(R, carpetas, datos)), ('G2', lambda: gateG2(R, carpetas, datos)),
        ('G3', lambda: gateG3(R, carpetas, datos, mes)), ('G4', lambda: gateG4(R, carpetas, datos)),
        ('G5', lambda: gateG5(R, carpetas, datos)), ('G6', lambda: gateG6(R, carpetas, datos, cifras, cifras_path)),
        ('G7', lambda: gateG7(R, carpetas, datos, cifras)), ('G8', lambda: gateG8(R, carpetas, datos, cifras)),
        ('G9', lambda: gateG9(R, carpetas, datos)),
    ])
    for g, fn in gates.items():
        if solo and g not in solo:
            continue
        if g == 'G9' and not g9:
            continue
        try:
            fn()
        except Exception as e:                                   # noqa: BLE001
            R.f(g, 'EXCEPCIÓN %r' % e)
        if not silencioso:
            fl = R.fallos.get(g, [])
            print('%s %s' % (g, 'VERDE' if not fl else 'ROJO (%d)' % len(fl)))
            for x in R.info.get(g, []):
                print('   · ' + x)
            for x in fl[:25]:
                print('   ✗ ' + x)
    return R


# ==========================================================================
# Autotest: defectos inyectados en una COPIA
# ==========================================================================
def _wb_mutar(base, corto, fn, recalcular=False):
    p = ruta_en(carpetas_de(base), corto)
    wb = openpyxl.load_workbook(p)
    fn(wb)
    wb.save(p)
    if recalcular:
        A.inyectar(p)


def _zip_props(base, corto, viejo, nuevo):
    p = ruta_en(carpetas_de(base), corto)
    with zipfile.ZipFile(p) as z:
        infos = z.infolist()
        partes = {i.filename: z.read(i.filename) for i in infos}
    x = partes['docProps/core.xml'].decode('utf-8')
    if viejo not in x:
        raise RuntimeError('docProps sin %r' % viejo)
    partes['docProps/core.xml'] = x.replace(viejo, nuevo).encode('utf-8')
    tmp = p + '.tmp'
    with zipfile.ZipFile(tmp, 'w', zipfile.ZIP_DEFLATED) as zout:
        for i in infos:
            zout.writestr(i, partes[i.filename])
    os.replace(tmp, p)


def _primera_cf_text(wb, hoja, texto):
    for cf in wb[hoja].conditional_formatting:
        for r in cf.rules:
            if r.text == texto:
                return r
    raise RuntimeError('sin CF %r' % texto)


def _cifras_mal(base):
    p = os.path.join(base, 'cifras_caso.json')
    c = json.load(open(p, encoding='utf-8'))
    c['ft']['ventas_a1']['texto'] = '$1'
    json.dump(c, open(p, 'w', encoding='utf-8'))


def _dv_lista(wb, hoja, nueva):
    for dv in wb[hoja].data_validations.dataValidation:
        if dv.type == 'list':
            dv.formula1 = nueva
            return
    raise RuntimeError('sin DV de lista en %s' % hoja)


def _tarea_fuera(wb):
    ws = wb['Phase 1 - Business Setup']
    for c in ws['A']:
        if c.value == '☐':
            c.value = None
            return
    raise RuntimeError('sin ☐')


def _desbloquear(wb):
    ws = wb['3-Year P&L']
    c = ws['A9']
    pr = copy.copy(c.protection)
    pr.locked = False
    c.protection = pr


def _e2_revertir(wb):
    ws = wb['3-Year P&L']
    ws['G13'].value = ws['G13'].value.replace("'0. Assumptions'!$B$41", "'0. Assumptions'!$B$39")


MUTACIONES = [
    ('texto en español en una celda', 'FTP', lambda b: _wb_mutar(b, 'FTP', lambda w: w['Startup Costs'].cell(
        row=3, column=1, value='Inversión total del vehículo')), {'G2'}),
    ('fórmula alterada', 'FTP', lambda b: _wb_mutar(b, 'FTP', lambda w: w['3-Year P&L'].__setitem__(
        'B9', '=IFERROR(IF(B6*B7*B8=0,"",B6*B7*B8*1.1),"")')), {'G1'}),
    ('hoja en A4', 'FTC', lambda b: _wb_mutar(b, 'FTC', lambda w: setattr(w['Phase 2 - Truck & Permits'].page_setup,
                                                                          'paperSize', 9)), {'G3'}),
    ('formato con €', 'FTP', lambda b: _wb_mutar(b, 'FTP', lambda w: setattr(w['3-Year P&L']['B9'], 'number_format',
                                                                             '#,##0 €')), {'G2'}),
    ('ítem de DV en español', 'FTP', lambda b: _wb_mutar(b, 'FTP', lambda w: _dv_lista(w, 'Startup Costs', '"Sí,No"')),
     {'G4'}),
    ('token de CF que colisiona', 'FTP', lambda b: _wb_mutar(b, 'FTP', lambda w: setattr(
        _primera_cf_text(w, 'Instructions', 'OK'), 'text', 'REVIEW')), {'G5'}),
    ('contraseña de protección', 'CAFP', lambda b: _wb_mutar(b, 'CAFP', lambda w: setattr(
        w['Staffing'].protection, 'password', 'x')), {'G1'}),
    ('celda desbloqueada fuera de verde', 'CAFP', lambda b: _wb_mutar(b, 'CAFP', _desbloquear), {'G1'}),
    ('aviso D21 borrado', 'CAFC', lambda b: _wb_mutar(b, 'CAFC', lambda w: w['Instructions'].__setitem__('A12', None)),
     {'G3'}),
    ('línea de versión vieja', 'CAFC', lambda b: _wb_mutar(b, 'CAFC', lambda w: w['Instructions'].__setitem__(
        'A11', mapas.version_en('caf', 'September 2026'))), {'G3'}),
    ('docProps en español', 'FTC', lambda b: _zip_props(b, 'FTC', 'Food Truck Startup Checklist (68 Tasks)',
                                                        'Checklist de apertura del food truck'), {'G2'}),
    ('«SBA minimum» en una nota', 'CAFP', lambda b: _wb_mutar(b, 'CAFP', lambda w: w['Financing'].__setitem__(
        'C17', 'A 1.25x DSCR is the SBA minimum for this loan')), {'G7'}),
    ('libro guardado sin caché', 'FTP', lambda b: _wb_mutar(b, 'FTP', lambda w: None), {'G6'}),
    ('caso que rompe D15', 'FTP', lambda b: _wb_mutar(b, 'FTP', lambda w: w['0. Assumptions'].__setitem__('B4', 55),
                                                      recalcular=True), {'G6', 'G7'}),
    ('regla US cambiada', 'CAFP', lambda b: _wb_mutar(b, 'CAFP', lambda w: w['0. Assumptions'].__setitem__('B39', 0.07),
                                                      recalcular=True), {'G7'}),
    ('merge añadido', 'FTP', lambda b: _wb_mutar(b, 'FTP', lambda w: w['Staffing'].merge_cells('A30:C30')), {'G1'}),
    ('tarea borrada del checklist', 'FTC', lambda b: _wb_mutar(b, 'FTC', _tarea_fuera), {'G1'}),
    ('E2 revertido', 'FTP', lambda b: _wb_mutar(b, 'FTP', _e2_revertir, recalcular=True), {'G1', 'G6'}),
    ('carácter no latino', 'CAFC', lambda b: _wb_mutar(b, 'CAFC', lambda w: w['Phase 5 - Marketing'].__setitem__(
        'B3', 'Social media plan 计划')), {'G2'}),
    ('cifras_caso.json desalineado', '-', _cifras_mal, {'G6'}),
]


def autotest(base, mes):
    origen = os.path.abspath(base) if base else None
    tmp = tempfile.mkdtemp(prefix='bp-autotest-')
    try:
        fuente = tmp + '/fuente'
        os.makedirs(fuente)
        for plan, c in carpetas_de(origen).items():
            shutil.copytree(c, os.path.join(fuente, mapas.PLANES[plan]['slug']))
        cif, _p = cargar_cifras(origen)
        json.dump(cif, open(os.path.join(fuente, 'cifras_caso.json'), 'w', encoding='utf-8'))
        print('== autotest: línea base (copia limpia)', flush=True)
        R0 = ejecutar(fuente, mes, solo={'G1', 'G2', 'G3', 'G4', 'G5', 'G6', 'G7'}, silencioso=True)
        if R0.gates_rojos():
            print('  la copia limpia NO está verde: %s' % R0.gates_rojos())
            for g in R0.gates_rojos():
                for x in R0.fallos[g][:5]:
                    print('   ✗ %s %s' % (g, x))
            return 1
        print('  copia limpia VERDE en G1-G7')
        cazados = 0
        for nombre, _corto, fn, esperado in MUTACIONES:
            d = os.path.join(tmp, 'm')
            shutil.rmtree(d, ignore_errors=True)
            shutil.copytree(fuente, d)
            fn(d)
            R = ejecutar(d, mes, solo={'G1', 'G2', 'G3', 'G4', 'G5', 'G6', 'G7'}, silencioso=True)
            rojos = set(R.gates_rojos())
            ok = esperado <= rojos
            cazados += ok
            print('  %s %-38s esperado %-10s rojos %s' % ('OK ' if ok else 'FALLA', nombre, ','.join(sorted(esperado)),
                                                         ','.join(sorted(rojos)) or '—'))
        print('autotest: %d/%d defectos cazados' % (cazados, len(MUTACIONES)))
        return 0 if cazados == len(MUTACIONES) else 1
    finally:
        shutil.rmtree(tmp, ignore_errors=True)


def main():
    ap = argparse.ArgumentParser(description='Gates F2 de los Business Plan Kits EN (G1-G9)')
    ap.add_argument('--dir', default=None, help='dry-run con subcarpetas <slug>/ y cifras_caso.json')
    ap.add_argument('--mes', default=mapas.MES_VERSION.split()[0])
    ap.add_argument('--solo', default=None, help='p. ej. G1,G2')
    ap.add_argument('--json', default=None)
    ap.add_argument('--autotest', action='store_true')
    args = ap.parse_args()
    mes = '%s %s' % (args.mes, mapas.MES_VERSION.split()[1])
    t0 = time.time()
    if args.autotest:
        sys.exit(autotest(args.dir, mes))
    solo = set(args.solo.split(',')) if args.solo else None
    R = ejecutar(os.path.abspath(args.dir) if args.dir else None, mes, solo)
    rojos = R.gates_rojos()
    if args.json:
        json.dump({'fallos': R.fallos, 'info': R.info}, open(args.json, 'w', encoding='utf-8'), ensure_ascii=False,
                  indent=1)
    print('\n%s (%.0f s)' % ('TODO VERDE' if not rojos else 'ROJO: ' + ', '.join(rojos), time.time() - t0))
    sys.exit(1 if rojos else 0)


if __name__ == '__main__':
    main()
