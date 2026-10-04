#!/usr/bin/env python3
"""
mapas.py — Food Truck Business Plan Kit + Coffee Shop Business Plan Kit (EN) · mapas ES→EN.

SPEC: `scripts/productos-digitales/business-plans-en/SPEC.md` (D1-D29, §1.1 E1-E2, §2-§5). Sesión Claude Code,
4-oct-2026. Calcado de `financial-kit/mapas.py`. Importable SIN efectos (constantes y funciones). Ejecutado como
script corre el autotest (pestañas ≤ 31 sin caracteres prohibidos, claves EN sin comillas ni español, límites de
Excel, reglas US con fuente, tokens del docx) y, si existe `censo_es.json` al lado, el cruce con el censo
(pestañas, ítems de DV, literales de fórmula/CF cubiertos por CLAVES o FORMULAS_EN, CF que no cambian de color,
VALORES_EN sobre celdas VERDES con su rótulo, tokens resueltos).

    python3 mapas.py

Contenido:
  · PLANES, FICHEROS, TITULOS, LIBROS_ES, DOCX         nombres, slugs y títulos (D1, D2, D5, §2.1)
  · HOJAS_POR_LIBRO, hoja_en()                          pestañas (D6, §2.2)
  · CLAVES                                              literales de fórmula/CF/DV e ítems de lista (D7)
  · FORMULAS_EN                                         fórmulas con texto, celda a celda, ya en EN (D7, D11)
  · PARCHES_FORMULA                                     E2 (SPEC §1.1), en espacio ES
  · VALORES_EN                                          caso de ejemplo US (D9-D16, research §7)
  · FORMATOS_EN, FOOTER_EN, version_en(), DOCPROPS_EN   D8, D20
  · AVISO_EN, TEXTOS_NUEVOS                             D21 / E1 (4 celdas)
  · FIJOS, POR_CELDA, CELDA_PISTAS, ZONAS_CELDA         textos que no traducen los subagentes / pistas por celda
  · REGLAS_US, GLOSARIO                                 SPEC §2.3 y §3 con fuente
  · TOKENS_DOCX, DOCX_FIJOS, DOCX_MOVER, ENCABEZADOS_EN docx (D22, D23) y `formatear()` / `calcular_cifras()`
  · restos_espanol(), no_latinos()                      detectores compartidos por los checkers
"""
import json
import math
import os
import re
import unicodedata
from collections import OrderedDict

AQUI = os.path.dirname(os.path.abspath(__file__))
# Conjunto de planes activo (F2 tanda 1 de Restaurant + Bakery, 4-oct-2026): por defecto food truck + cafetería con
# sus datos en esta carpeta; con BP_CONJUNTO=restpan, `../business-plans-en-2/mapas_rp.py` sustituye al final de
# este módulo los datos por los de restaurante + panadería (DATOS = esa carpeta). El código de los scripts es el
# mismo para los dos conjuntos: lee los datos de `mapas.DATOS` y los nombres de `mapas.*`.
CONJUNTO = os.environ.get('BP_CONJUNTO', 'ftcaf')
DATOS = AQUI

# ==========================================================================
# 0. Productos (D1-D3, D20)
# ==========================================================================
PLANES = OrderedDict([
    ('ft', OrderedDict([
        ('slug', 'food-truck-business-plan'),
        ('producto', 'Food Truck Business Plan Kit'),
        ('slug_es', 'plan-negocio-food-truck'),
        ('dir_es', 'astro-site/public/dl/plan-negocio-food-truck'),
        ('env_stripe', 'VITE_STRIPE_PAYMENT_LINK_FOOD_TRUCK_BUSINESS_PLAN'),
        ('keywords', 'food truck business plan, food truck business plan template, food truck startup costs, '
                     'food truck permits, how to start a food truck business, AI Chef Pro'),
    ])),
    ('caf', OrderedDict([
        ('slug', 'coffee-shop-business-plan'),
        ('producto', 'Coffee Shop Business Plan Kit'),
        ('slug_es', 'plan-negocio-cafeteria'),
        ('dir_es', 'astro-site/public/dl/plan-negocio-cafeteria'),
        ('env_stripe', 'VITE_STRIPE_PAYMENT_LINK_COFFEE_SHOP_BUSINESS_PLAN'),
        ('keywords', 'coffee shop business plan, cafe business plan, coffee shop business plan template, '
                     'coffee shop startup costs, how to open a coffee shop, AI Chef Pro'),
    ])),
])
for _p in PLANES.values():
    _p['url'] = 'aichef.pro/en/digital-products/' + _p['slug']
    _p['dir_en'] = 'astro-site/public/dl/' + _p['slug']
PRECIO = '$39'
VERSION = '2.2'
URL_TIENDA = 'aichef.pro/en/digital-products'
FOOTER_ES = 'AI Chef Pro · aichef.pro · Página &P de &N'
FOOTER_EN = 'AI Chef Pro · aichef.pro · Page &P of &N'

# ==========================================================================
# 1. Ficheros (D5, SPEC §2.1). corto: FTP/FTC = plan financiero / checklist food truck; CAFP/CAFC = cafetería
# ==========================================================================
LIBROS = OrderedDict([
    # corto: (plan, fichero ES, fichero EN, título EN)
    ('FTP', ('ft', 'plan-financiero-food-truck.xlsx', 'food-truck-financial-projections.xlsx',
             'Food Truck Financial Projections: 3-Year P&L, Cash Flow & Financing')),
    ('FTC', ('ft', 'checklist-apertura-food-truck.xlsx', 'food-truck-startup-checklist.xlsx',
             'Food Truck Startup Checklist (68 Tasks)')),
    ('CAFP', ('caf', 'plan-financiero-cafeteria-brunch.xlsx', 'coffee-shop-financial-projections.xlsx',
              'Coffee Shop Financial Projections: 3-Year P&L, Cash Flow & Financing')),
    ('CAFC', ('caf', 'checklist-apertura-cafeteria-brunch.xlsx', 'coffee-shop-opening-checklist.xlsx',
              'Coffee Shop Opening Checklist (75 Tasks)')),
])
DOCX = OrderedDict([
    ('ft', ('plan-de-negocio-food-truck.docx', 'food-truck-business-plan.docx', 'Food Truck Business Plan (10 Sections)')),
    ('caf', ('plan-de-negocio-cafeteria-brunch.docx', 'coffee-shop-business-plan.docx',
             'Coffee Shop Business Plan (10 Sections)')),
])
FICHEROS = OrderedDict((v[1], v[2]) for v in LIBROS.values())
FICHEROS.update((v[0], v[1]) for v in DOCX.values())
TITULOS = OrderedDict((v[2], v[3]) for v in LIBROS.values())
TITULOS.update((v[1], v[2]) for v in DOCX.values())
LIBROS_ES = [v[1] for v in LIBROS.values()]
CORTO_DE = {v[1]: k for k, v in LIBROS.items()}
PLAN_DE = {k: v[0] for k, v in LIBROS.items()}
TIPO_DE = {'FTP': 'plan', 'CAFP': 'plan', 'FTC': 'check', 'CAFC': 'check'}


def corto_de(fichero):
    return CORTO_DE[fichero]


def ruta_es(corto, repo):
    plan, f_es = LIBROS[corto][0], LIBROS[corto][1]
    return os.path.join(repo, PLANES[plan]['dir_es'], f_es)


def cortos_de(plan, tipo=None):
    return [c for c, v in LIBROS.items() if v[0] == plan and (tipo is None or TIPO_DE[c] == tipo)]


def corto_plan(plan):
    """Libro del plan financiero de un plan (FTP, CAFP, RESTP, PANP)."""
    return cortos_de(plan, 'plan')[0]


def corto_check(plan):
    return cortos_de(plan, 'check')[0]


# Grupos de traducción (textos_es.json → textos_en/<grupo>.json). GX = no se traduce.
GRUPOS = ('GM', 'GFT', 'GCAF')                     # los que lee aplicar_en.py
GRUPOS_TRADUCTORES = GRUPOS                        # los que van a los subagentes (preparar_tandas / check_textos)
GRUPOS_OVERRIDES = ('GFT', 'GCAF')                 # textos_en/overrides_<g>.json (EN por celda)
GRUPOS_DESC = OrderedDict([
    ('GM', 'Comunes a FT y CAF (motor 2.2 + molde del checklist): se traducen PRIMERO y fijan el glosario'),
    ('GFT', 'Solo food truck (plan financiero + checklist FT), incluidas las celdas de las tablas D18/D19'),
    ('GCAF', 'Solo cafetería (plan financiero + checklist CAF), incluidas las celdas de las tablas D18/D19'),
    ('GX', 'NO se traduce: pestañas (mapas.HOJAS_POR_LIBRO), CLAVES, FIJOS, POR_CELDA, versión, pie y docProps'),
])


def grupo_de(planes, texto=None, tratamiento=None):
    """Grupo de una cadena traducible por los planes en los que aparece (lista ordenada)."""
    return 'GM' if planes == ['caf', 'ft'] else ('GFT' if planes == ['ft'] else 'GCAF')


# ==========================================================================
# 2. Pestañas (D6, SPEC §2.2) — ≤ 31 caracteres, sin []:*?/\
# ==========================================================================
_HOJAS_PLAN = OrderedDict([
    ('0. Supuestos', '0. Assumptions'), ('Inversión Inicial', 'Startup Costs'), ('PyG 3 Años', '3-Year P&L'),
    ('Punto Equilibrio', 'Break-Even'), ('Escenarios', 'Scenarios'), ('Personal', 'Staffing'),
    ('Instrucciones', 'Instructions'), ('Tesorería 12 meses', '12-Month Cash Flow'), ('Financiación', 'Financing'),
])
_HOJAS_CHECK = OrderedDict([
    ('F1 - Constitución', 'Phase 1 - Business Setup'), ('F2 - Vehículo', 'Phase 2 - Truck & Permits'),
    ('F2 - Local', 'Phase 2 - Location & Permits'), ('F3 - Equipamiento', 'Phase 3 - Equipment'),
    ('F4 - Personal', 'Phase 4 - Staff'), ('F5 - Marketing', 'Phase 5 - Marketing'),
    ('F6 - 90 Días', 'Phase 6 - First 90 Days'), ('Instrucciones', 'Instructions'),
])
HOJAS_POR_LIBRO = OrderedDict([
    ('FTP', _HOJAS_PLAN), ('CAFP', _HOJAS_PLAN),
    ('FTC', OrderedDict((k, v) for k, v in _HOJAS_CHECK.items() if k != 'F2 - Local')),
    ('CAFC', OrderedDict((k, ('Phase 3 - Build-Out & Equipment' if k == 'F3 - Equipamiento' else v))
                         for k, v in _HOJAS_CHECK.items() if k != 'F2 - Vehículo')),
])
HOJAS_ES_POR_LIBRO = OrderedDict([
    ('FTP', ['0. Supuestos', 'Inversión Inicial', 'PyG 3 Años', 'Punto Equilibrio', 'Escenarios', 'Personal',
             'Instrucciones', 'Tesorería 12 meses', 'Financiación']),
    ('FTC', ['F1 - Constitución', 'F2 - Vehículo', 'F3 - Equipamiento', 'F4 - Personal', 'F5 - Marketing',
             'F6 - 90 Días', 'Instrucciones']),
])
HOJAS_ES_POR_LIBRO['CAFP'] = list(HOJAS_ES_POR_LIBRO['FTP'])
HOJAS_ES_POR_LIBRO['CAFC'] = [h.replace('F2 - Vehículo', 'F2 - Local') for h in HOJAS_ES_POR_LIBRO['FTC']]
HOJA_PROHIBIDOS = set('[]:*?/\\')
# Tareas por fase (D24): FT 11/13/12/10/12/10 = 68 · CAF 11/13/15/12/14/10 = 75
TAREAS_POR_FASE = OrderedDict([('FTC', [11, 13, 12, 10, 12, 10]), ('CAFC', [11, 13, 15, 12, 14, 10])])


def norm_hoja(h):
    """Pestaña del molde numerado del bar-restaurante («1. Inversión Inicial», «2. P&L 3 Años») → nombre canónico
    (el de FT/CAF, que es el que usan TOKENS_DOCX y el código). En FT/CAF es la identidad."""
    if h == '0. Supuestos':
        return h
    h = re.sub(r'^\d+\. ', '', h)
    return {'P&L 3 Años': 'PyG 3 Años'}.get(h, h)


def hoja_es(corto, h):
    """Nombre REAL de la pestaña ES del libro para un nombre real o canónico."""
    hojas = HOJAS_POR_LIBRO[corto]
    if h in hojas:
        return h
    for real in hojas:
        if norm_hoja(real) == h:
            return real
    raise KeyError('%s: pestaña %r' % (corto, h))


def hoja_en(corto, hoja_es_):
    return HOJAS_POR_LIBRO[corto][hoja_es(corto, hoja_es_)]


# Molde de cada checklist: 'hojas' = una hoja por fase, OK en la col. A (☐/✓/N/A); 'cabeceras' = una sola hoja con
# las fases como filas-cabecera fusionadas (RESTC), ver FASES_CABECERA
MOLDE_CHECK = OrderedDict([('FTC', 'hojas'), ('CAFC', 'hojas')])
COL_OK = OrderedDict([('FTC', 1), ('CAFC', 1)])            # columna de la DV ✓ de cada checklist
FASES_CABECERA = OrderedDict()                              # corto → (hoja ES, [filas-cabecera], última fila de tarea)
VALORES_OK = ('☐', '✓', 'N/A')


def hojas_fase(corto):
    return [h for h in HOJAS_ES_POR_LIBRO[corto] if h != 'Instrucciones']


def contar_tareas(corto, wb, en=True):
    """Tareas por fase de un checklist abierto (EN si en=True; ES si no). 'hojas': ☐/✓/N/A de la col. A de cada hoja
    de fase; 'cabeceras': filas con texto en la col. B entre una fila-cabecera y la siguiente."""
    nombre = (lambda h: hoja_en(corto, h)) if en else (lambda h: h)
    if MOLDE_CHECK.get(corto, 'hojas') == 'hojas':
        return [sum(1 for c in wb[nombre(h)]['A'] if c.value in VALORES_OK) for h in hojas_fase(corto)]
    h, cabs, fin = FASES_CABECERA[corto]
    ws = wb[nombre(h)]
    lims = list(cabs) + [fin + 1]
    return [sum(1 for r in range(lims[k] + 1, lims[k + 1]) if isinstance(ws.cell(r, 2).value, str) and
                ws.cell(r, 2).value.strip()) for k in range(len(cabs))]


def todas_hojas_es():
    return {h for d in HOJAS_POR_LIBRO.values() for h in d}


# ==========================================================================
# 3. Claves (D7): literales de fórmula, de CF y de DV, e ítems de lista. Un solo paso ES→EN. ✓ ☐ N/A se quedan.
# ==========================================================================
CLAVES = OrderedDict([
    ('Sí', 'Yes'), ('No', 'No'),                    # DV Inversión col. E, SUMIF, «Punto de equilibrio alcanzado»
    ('CUMPLE', 'OK'), ('REVISAR', 'REVIEW'),        # Instrucciones: ratios auditados (CF containsText)
    ('Más de 3 años', 'Over 3 years'),              # Tesorería: payback
    ('✓', '✓'), ('☐', '☐'), ('N/A', 'N/A'), ('—', '—'),
])

# ==========================================================================
# 4. Fórmulas que cambian. FORMULAS_EN: (corto, hoja ES, celda o rango de UNA fila) → fórmula EN FINAL (pestañas EN
#    y literales EN; «{c}» = letra de columna de la celda si la clave es un rango). Cubre las fórmulas con literales
#    de texto (D7) y Personal!A2 (D11). La rellena `_formulas_en()` más abajo.
#    PARCHES_FORMULA (E2, SPEC §1.1): (corto, hoja ES, celda) → (viejo, nuevo) en ESPACIO ES.
# ==========================================================================
FORMULAS_EN = OrderedDict()
PARCHES_FORMULA = OrderedDict([
    (('FTP', 'PyG 3 Años', 'G13'), ("'0. Supuestos'!$B$39", "'0. Supuestos'!$B$41")),
    (('FTP', 'PyG 3 Años', 'G14'), ("('0. Supuestos'!$B$62*'0. Supuestos'!$B$40+(1-'0. Supuestos'!$B$62)*"
                                    "'0. Supuestos'!$B$39)", "'0. Supuestos'!$B$41")),
    (('CAFP', 'PyG 3 Años', 'G13'), ("'0. Supuestos'!$B$39", "'0. Supuestos'!$B$41")),
    (('CAFP', 'PyG 3 Años', 'G14'), ("('0. Supuestos'!$B$62*'0. Supuestos'!$B$40+(1-'0. Supuestos'!$B$62)*"
                                     "'0. Supuestos'!$B$39)", "'0. Supuestos'!$B$41")),
])


def celdas_e2(corto):
    """(hoja ES, celda) del E2: soportado de compras → '0. Supuestos'!$B$41 (= 0 en la EN)."""
    return [(h, x) for (c, h, x), (_v, n) in PARCHES_FORMULA.items() if c == corto and n.endswith('$B$41')]


def parchear_formula_es(corto, hoja_es, coord, f):
    par = PARCHES_FORMULA.get((corto, hoja_es, coord))
    if par is None:
        return f
    if par[0] not in f:
        raise ValueError('PARCHES_FORMULA %s %s!%s: %r no está en %r' % (corto, hoja_es, coord, par[0], f[:140]))
    return f.replace(par[0], par[1])


# ==========================================================================
# 5. Números que cambian (D9-D16; research §7). (corto, hoja ES, celda) → (valor, rótulo ES de la fila, col. A,
#    por el que EMPIEZA). None = celda vacía. 'CALIBRAR' = lo fija aplicar_en (préstamo = el que propone
#    `Financiación`, redondeado a 1.000 por arriba). Toda celda debe ser VERDE en el ES (gate_f1).
# ==========================================================================
CALIBRAR = 'CALIBRAR'


def _v(corto, hoja, filas_valores):
    return [((corto, hoja, c), v) for c, v in filas_valores]


_SUP_COMUN = [  # (celda, (valor FT, valor CAF), rótulo)
    ('B4', (70, 140), 'Clientes/día'), ('B5', (14, 10.5), 'Ticket medio'), ('B6', (260, 340), 'Días de apertura'),
    ('B13', (0.30, 0.30), 'Coste de mercancía sobre las ventas de COMIDA'),
    ('B14', (0.25, 0.24), 'Coste de mercancía sobre las ventas de BEBIDA'),
    ('B15', (0.04, 0.03), 'Consumibles'), ('B16', (0, 0), 'Ventas por delivery'),
    ('B17', (0.25, 0.25), 'Comisión de la plataforma'), ('B18', (0.035, 0.035), 'Comisión de los medios de pago'),
    ('B20', (0.10, 0.10), 'Seguridad Social'), ('B21', (12, 12), 'Número de pagas'),
    ('B22', (15080, 15080), 'SMI anual'),
    ('B24', (1200, 4500), 'Alquiler mensual'), ('B25', (1, 2), 'Fianza'), ('B26', (500, 1200), 'Suministros'),
    ('B27', (4000, 3500), 'Seguros'),
    ('B30', (35000, 60000), 'Recursos propios'), ('B31', (CALIBRAR, CALIBRAR), 'Préstamo bancario'),
    ('B32', (0.10, 0.10), 'Tipo de interés'), ('B33', (7, 10), 'Plazo del préstamo'),
    ('B34', (0, 0), 'Carencia'),
    ('B37', (0.25, 0.25), 'Impuesto de Sociedades, nueva'), ('B38', (0.25, 0.25), 'Impuesto de Sociedades, tipo'),
    ('B39', (0.08, 0.08), 'IVA repercutido de restauración'), ('B40', (0.08, 0.08), 'IVA repercutido/soportado'),
    ('B41', (0, 0), 'IVA soportado en compras'), ('B42', (0, 0), 'Bases negativas'),
    ('B44', (7, 10), 'Vida útil de obra'), ('B45', (7, 7), 'Vida útil de maquinaria'),
    # B58: el ES trae 0, que viola su propia DV (1-365). US: el abono de tarjeta llega el día hábil siguiente (F2-NOTAS §6)
    ('B58', (1, 1), 'Días medios de cobro'),
    ('B59', (7, 14), 'Días medios de pago'), ('B62', (0, 0), 'Bebida ALCOHÓLICA'),
    ('B63', (0.08, 0.08), 'IVA de la bebida ALCOHÓLICA'),
    ('B65', (1.15, 1.15), 'DSCR mínimo'), ('B66', (1.25, 1.25), 'DSCR objetivo'),
    ('B68', (None, None), 'Suelo salarial anual'),
]
VALORES_EN = OrderedDict()
for _k, (_c, (_vft, _vcaf), _rot) in enumerate(_SUP_COMUN):
    VALORES_EN[('FTP', '0. Supuestos', _c)] = (_vft, _rot)
    VALORES_EN[('CAFP', '0. Supuestos', _c)] = (_vcaf, _rot)
# Inversión: partidas US en las MISMAS filas (research §7)
for _i, (_val, _rot) in enumerate([
        (45000, 'Vehículo'), (20000, 'Adaptación'), (15000, 'Equipamiento cocina'), (5000, 'Generador'),
        (4000, 'Instalación gas'), (1500, 'Depósito de agua'), (800, 'Constitución'), (3000, 'Licencias'),
        (1000, 'Proyecto técnico'), (1000, 'TPV'), (2000, 'Menaje'), (2000, 'Stock'), (3000, 'Marketing'),
        (1500, 'Toldo')]):
    VALORES_EN[('FTP', 'Inversión Inicial', 'B%d' % (7 + _i))] = (_val, _rot)
for _i, (_val, _rot) in enumerate([
        (1500, 'Constitución'), (2000, 'Licencia de actividad'), (8000, 'Proyecto técnico'), (40000, 'Obra civil'),
        (15000, 'Instalación eléctrica'), (10000, 'Fontanería'), (3000, 'Molinillo'), (6000, 'Horno'),
        (5000, 'Vitrina'), (6000, 'Neveras'), (2000, 'Plancha'), (1500, 'Batidora'), (5000, 'Lavavajillas'),
        (10000, 'Mobiliario sala'), (12000, 'Barra'), (5000, 'Terraza'), (3000, 'Vajilla'), (6000, 'Decoración'),
        (4000, 'Rotulación'), (4000, 'Marketing'), (5000, 'Stock'), (15000, 'Máquina de café'), (3000, 'TPV')]):
    VALORES_EN[('CAFP', 'Inversión Inicial', 'B%d' % (7 + _i))] = (_val, _rot)
# P&L: fijos US (mismas filas) y umbrales D16 [kit benchmark]
for _i, (_val, _rot) in enumerate([
        (2000, 'Permisos'), (3000, 'Mantenimiento'), (6000, 'Marketing'), (2400, 'Gestoría'), (1200, 'Software'),
        (1200, 'Limpieza'), (600, 'Gestión de residuos'), (400, 'Desinsectación'), (300, 'Prevención'),
        (200, 'Derechos de autor')]):
    VALORES_EN[('FTP', 'PyG 3 Años', 'B%d' % (26 + _i))] = (_val, _rot)
for _i, (_val, _rot) in enumerate([
        (3000, 'Gestoría'), (2000, 'Mantenimiento'), (6000, 'Marketing'), (1800, 'Software'), (2400, 'Limpieza'),
        (1200, 'Gestión de residuos'), (800, 'Desinsectación'), (600, 'Derechos de autor'), (500, 'Prevención')]):
    VALORES_EN[('CAFP', 'PyG 3 Años', 'B%d' % (26 + _i))] = (_val, _rot)
VALORES_EN.update([
    (('FTP', 'PyG 3 Años', 'E51'), (0.58, 'Margen bruto')), (('FTP', 'PyG 3 Años', 'E53'), (0.35, 'Coste de mercancía /')),
    (('FTP', 'PyG 3 Años', 'E54'), (0.30, 'Coste de personal')), (('FTP', 'PyG 3 Años', 'E55'), (0.06, 'Aparcamiento')),
    (('FTP', 'PyG 3 Años', 'E56'), (0.06, 'Resultado neto')),
    # CAF E50 (margen bruto 65 %): research §7 no da umbral → se conserva el del ES [kit benchmark] (F2-NOTAS §3)
    (('CAFP', 'PyG 3 Años', 'E52'), (0.35, 'Coste de mercancía /')),
    (('CAFP', 'PyG 3 Años', 'E53'), (0.35, 'Coste de personal')), (('CAFP', 'PyG 3 Años', 'E54'), (0.12, 'Alquiler')),
    (('CAFP', 'PyG 3 Años', 'E55'), (0.06, 'Resultado neto')),
])
# Personal (D11; research §7): jornada (C) y bruto mensual total del puesto (D); 40 h; 50 semanas; horas y presencia
for _c, (_j, _b) in zip(range(5, 9), [(1.0, 3000), (0.8, 2218), (0.15, 390), (0.06, 166)]):
    VALORES_EN[('FTP', 'Personal', 'C%d' % _c)] = (_j, None)
    VALORES_EN[('FTP', 'Personal', 'D%d' % _c)] = (_b, None)
for _c, (_j, _b) in zip(range(5, 11), [(1.0, 3500), (1.0, 2773), (1.0, 2600), (0.8, 2496), (0.45, 1170), (0.1, 260)]):
    VALORES_EN[('CAFP', 'Personal', 'C%d' % _c)] = (_j, None)
    VALORES_EN[('CAFP', 'Personal', 'D%d' % _c)] = (_b, None)
VALORES_EN.update([
    (('FTP', 'Personal', 'B16'), (40, 'Jornada completa')), (('FTP', 'Personal', 'B17'), (50, 'Semanas')),
    (('FTP', 'Personal', 'B18'), (10, 'Horas de servicio')), (('FTP', 'Personal', 'B19'), (1.5, 'Personas necesarias')),
    (('CAFP', 'Personal', 'B18'), (40, 'Jornada completa')), (('CAFP', 'Personal', 'B19'), (50, 'Semanas')),
    # CAF horas/día (13) y presencia (2) se conservan del ES: cobertura 8,700 / 8,840 = 98.4 % (D15 ≥ 98 %)
])
# Escenarios [kit estimate]: mismas proporciones que el ES sobre el caso base US (FT 30/45/65 · 10/12/14 · 250;
# CAF 80/100/115 · 9.20/9.80/10.40 · 290/300/310)
VALORES_EN.update([
    (('FTP', 'Escenarios', 'B5'), (47, 'Clientes/día')), (('FTP', 'Escenarios', 'D5'), (101, 'Clientes/día')),
    (('FTP', 'Escenarios', 'B6'), (11.67, 'Ticket')), (('FTP', 'Escenarios', 'D6'), (16.33, 'Ticket')),
    (('FTP', 'Escenarios', 'B7'), (260, 'Días')), (('FTP', 'Escenarios', 'D7'), (260, 'Días')),
    (('CAFP', 'Escenarios', 'B5'), (112, 'Clientes/día')), (('CAFP', 'Escenarios', 'D5'), (161, 'Clientes/día')),
    (('CAFP', 'Escenarios', 'B6'), (9.86, 'Ticket')), (('CAFP', 'Escenarios', 'D6'), (11.14, 'Ticket')),
    (('CAFP', 'Escenarios', 'B7'), (329, 'Días')), (('CAFP', 'Escenarios', 'D7'), (351, 'Días')),
])

# Calibración D15 (F2 tanda 2): celdas VERDES del caso que se apartan de research §7 para que el caso base cumpla
# D15 con un caso creíble para un prestamista US. (corto, hoja ES, celda) → (valor, motivo). Cada entrada está
# explicada en F2-NOTAS.md §6 y gates_en.py G7 lo exige (la celda debe aparecer allí).
CALIBRACION_D15 = OrderedDict([
    # FT: con 70 clientes/día la holgura sobre el equilibrio de caja salía 6 % (< 15 %) y el neto del año 1, 1.8 %.
    # 80/día × $14 × 260 = $291,200: dentro del rango $250-500K y por debajo de la media de $346K (research §4)
    (('FTP', '0. Supuestos', 'B4'), (80, 'D15: holgura de caja 6 % → ≥ 15 %; ventas $291K < media US $346K')),
    (('FTP', 'Escenarios', 'B5'), (53, 'mismas proporciones que el ES (30/45) sobre los 80 del caso calibrado')),
    (('FTP', 'Escenarios', 'D5'), (116, 'mismas proporciones que el ES (65/45) sobre los 80 del caso calibrado')),
    # CAF: con $10.50 la holgura salía 12 % y el neto del año 1, 4.3 %. $11.00 de ticket medio de cafetería CON brunch
    # (rango de transacción $7.81-$11.11, research §4; el brunch completo va a $14-20 en la tabla D19)
    (('CAFP', '0. Supuestos', 'B5'), (11.0, 'D15: holgura de caja 12 % → ≥ 15 %; ticket dentro de $7.81-$11.11')),
    (('CAFP', 'Escenarios', 'B6'), (10.33, 'mismas proporciones que el ES (9.20/9.80) sobre los $11.00 calibrados')),
    (('CAFP', 'Escenarios', 'D6'), (11.67, 'mismas proporciones que el ES (10.40/9.80) sobre los $11.00 calibrados')),
    # Ronda de arreglos (M1): con la comisión de tarjeta al 3.5 % combinado, el neto del año 1 de la CAF bajaba al 5.9 %
    # (semáforo REVIEW). 145 clientes/día, dentro del rango 100-250 de research §4; el ticket se queda en $11.00
    (('CAFP', '0. Supuestos', 'B4'), (145, 'M1: neto año 1 5.9 % → ≥ 6 % con la tarjeta al 3.5 %; research 100-250/día')),
    (('CAFP', 'Escenarios', 'B5'), (116, 'mismas proporciones que el ES (80/100) sobre los 145 del caso calibrado')),
    (('CAFP', 'Escenarios', 'D5'), (167, 'mismas proporciones que el ES (115/100) sobre los 145 del caso calibrado')),
])

# ==========================================================================
# 6. Formato (D8, D20, D21, E1)
# ==========================================================================
FORMATOS_EN = OrderedDict([('#,##0.00 €', '#,##0.00'), ('#,##0 €', '#,##0'), ('€#,##0.00', '#,##0.00'),
                           ('#,##0.00\\ €', '#,##0.00'), ('#,##0\\ €', '#,##0')])
PAPEL_EN = 1                                         # US Letter en las 32 hojas
MES_VERSION = 'October 2026'


def version_en(plan, mes=MES_VERSION):
    return 'Version %s · %s · %s · info@aichef.pro' % (VERSION, mes, PLANES[plan]['url'])


AVISO_EN = ('Planning tool, not financial, tax or legal advice. Requirements vary by state, county and city — check '
            'with your accountant and local authorities.')
AVISO_CHECK_EN = 'This checklist is a starting point, not an exhaustive list.'
FILA_VERSION = OrderedDict([('FTP', 73), ('CAFP', 75), ('FTC', 11), ('CAFC', 11)])     # Instrucciones!A<n>
TEXTOS_NUEVOS = OrderedDict()                         # E1: (corto, hoja ES, celda) → (texto, celda de la que copia estilo)
for _c, _f in FILA_VERSION.items():
    _t = AVISO_EN if TIPO_DE[_c] == 'plan' else AVISO_EN + ' ' + AVISO_CHECK_EN
    TEXTOS_NUEVOS[(_c, 'Instrucciones', 'A%d' % (_f + 1))] = (_t, 'A%d' % _f)


def docprops_en(corto_o_plan):
    """docProps D20: title = título EN · producto; subject = producto · v2.2; description = URL."""
    if corto_o_plan in LIBROS:
        plan, _fes, fen, tit = LIBROS[corto_o_plan]
    else:
        plan = corto_o_plan
        _fes, fen, tit = DOCX[plan]
    p = PLANES[plan]
    return OrderedDict([
        ('title', '%s · %s' % (tit, p['producto'])), ('subject', '%s · v%s' % (p['producto'], VERSION)),
        ('keywords', p['keywords']), ('description', p['url']), ('category', 'AI Chef Pro · Digital products'),
        # creator = «AI Chef Pro»: censo-entregables.py --fail lo exige (Fase A); la autoría de John va en la firma
        ('creator', 'AI Chef Pro'), ('lastModifiedBy', 'AI Chef Pro'),
    ])


# ==========================================================================
# 7. Textos que NO van a los subagentes
#    FIJOS: por TEXTO ES (cualquier celda de los 4 libros) → EN fijado por la SPEC.
#    POR_CELDA: (corto, hoja ES, celda) → EN (cuando el mismo ES va a dos EN o el texto depende del plan).
# ==========================================================================
FIRMA_EN = ('Designed by John Guerrero — chef and restaurant consultant since 2010, in the kitchen since age 17 · '
            'johnguerrero.es')
FIJOS = OrderedDict([
    # ---- Instrucciones (D20) — mismas cadenas que la tienda EN
    ('Más plantillas y kits del catálogo en aichef.pro/productos-digitales',
     'More templates and kits in the catalog at ' + URL_TIENDA),
    ('Diseñado por John Guerrero — chef y consultor gastronómico desde 2010, en cocina desde los 17 años · '
     'johnguerrero.es', FIRMA_EN),
    ('Para editar una celda que no esté en verde: Revisar → Desproteger hoja (no tiene contraseña).',
     'To edit a cell that is not green: Review → Unprotect Sheet (there is no password).'),
    # ---- Checklist Instrucciones A3 (D5, I5): nombre comercial, no el id interno
    ('Este fichero forma parte del producto «plan-negocio-food-truck» de AI Chef Pro.',
     'This file is part of the Food Truck Business Plan Kit by AI Chef Pro.'),
    ('Este fichero forma parte del producto «plan-negocio-cafeteria» de AI Chef Pro.',
     'This file is part of the Coffee Shop Business Plan Kit by AI Chef Pro.'),
    # ---- D18 / D19: cabeceras de las dos tablas de Instrucciones
    ('QUÉ HA CAMBIADO RESPECTO DE LA VERSIÓN 1.1', 'WHY THIS MODEL IS BUILT THIS WAY'),
    ('v1.1', 'Common shortcut'), ('v2.2', 'This workbook'), ('Por qué', 'Why'),
    ('Fuente', 'Source'), ('DATOS DE REFERENCIA DEL SECTOR', 'INDUSTRY BENCHMARKS'),
    # ---- 0. Supuestos: rótulos con decisión de mercado (D9-D11, D13)
    ('Ticket medio SIN IVA (€)', 'Average check (excl. sales tax)'),
    ('PVP equivalente con IVA (calculado)', 'Menu price incl. sales tax (calculated)'),
    ('Seguridad Social a cargo de la empresa', 'Employer payroll taxes'),
    # m8 (revisión final): en EE. UU. «pay periods» es la frecuencia de nómina; aquí son las pagas MENSUALES del modelo
    ('Número de pagas anuales', 'Monthly pay periods per year (keep 12)'),
    ('SMI anual de referencia (€)', 'Federal minimum wage, annual full-time'),
    ('Suelo salarial anual del convenio provincial (€)', 'State / local minimum wage, annual full-time'),
    ('Impuesto de Sociedades, nueva creación', 'Effective income tax rate (years 1-2)'),
    ('Impuesto de Sociedades, tipo general', 'Effective income tax rate (from year 3)'),
    ('IVA repercutido de restauración', 'Sales tax rate on sales'),
    ('IVA repercutido/soportado general', 'Sales tax rate on alcohol to go'),
    ('IVA de la bebida ALCOHÓLICA para consumir en el acto', 'Sales tax rate on alcohol served on site'),
    ('IVA de la bebida ALCOHÓLICA servida en sala', 'Sales tax rate on alcohol served on site'),
    ('IVA soportado en compras e inversión', 'Recoverable input tax on purchases and capex'),
    ('Bases negativas de ejercicios anteriores (€)', 'Net operating losses carried forward'),
    ('Carencia de principal (años)', 'Interest-only period (years)'),
    ('DSCR mínimo aceptable', 'DSCR limit (minimum acceptable)'),
    ('DSCR objetivo (verde)', 'DSCR target (green)'),
    # ---- Financiación filas 8-11 (D13). B1 (revisión final): el cuadro, los intereses del P&L y el DSCR leen SOLO la
    #      fila 7 (= '0. Assumptions'!B31); las filas 8-9 suman como fuentes pero no se amortizan, y lo dicen
    ('Línea ICO (avalada por el ICO, la concede tu banco)',
     'Other loan (not in the schedule or the DSCR — enter your loan in row 7)'),
    ('Préstamo participativo ENISA', 'SBA Microloan (not in the schedule or the DSCR — enter your loan in row 7)'),
    ('Business angels o socios inversores', 'Investors / partners (equity: dilutes, not repaid)'),
    ('Subvenciones autonómicas o locales', 'Local grants (usually paid after you spend)'),
    # ---- Tesorería filas 16-20 (D9, SPEC §4.1) y cabeceras «Mes n (€)» (D8)
    ('IVA repercutido del mes (memoria)', 'Sales tax collected (memo)'),
    ('IVA soportado del mes (memoria)', 'Recoverable input tax (memo)'),
    ('IVA a compensar arrastrado (inversión incluida)', 'Carried forward (capex input tax included)'),
    ('Resultado de la liquidación trimestral', 'Quarterly sales tax return'),
    ('Pago del IVA (modelo 303)', 'Sales tax remitted (quarterly filer)'),
    ('Año (€)', 'Year'),
    # ---- rótulos que CITAN las fórmulas de texto de FORMULAS_EN (Punto Equilibrio A26/A28): deben casar
    ('Ingresos necesarios al año', 'Revenue needed per year'),
    ('Ingresos necesarios al año (caja)', 'Revenue needed per year (cash)'),
])
FIJOS.update(('Mes %d (€)' % _m, 'Month %d' % _m) for _m in range(1, 13))
POR_CELDA = OrderedDict()

# Zonas que se traducen CELDA A CELDA (un id por celda, incluidas las celdas sin letras): el texto EN depende de la
# fila. (corto, hoja ES) → (fila1, fila2, columnas, motivo)
ZONAS_CELDA = OrderedDict([
    (('FTP', 'Instrucciones'), [(35, 51, 'ABCD', 'D18'), (55, 68, 'ABCD', 'D19')]),
    (('CAFP', 'Instrucciones'), [(None, None, 'ABCD', 'D18'), (None, None, 'ABCD', 'D19')]),
])
# las filas de la CAF se resuelven en extraer_textos.py por sus cabeceras (QUÉ HA CAMBIADO… / DATOS DE REFERENCIA…)
CABECERA_ZONA = OrderedDict([('D18', 'QUÉ HA CAMBIADO RESPECTO DE LA VERSIÓN 1.1'),
                             ('D19', 'DATOS DE REFERENCIA DEL SECTOR')])

PISTA_D18 = ('D18: tabla «WHY THIS MODEL IS BUILT THIS WAY» (columnas Topic | Common shortcut | This workbook | '
             'Why). Col. A = tema de la fila; col. B = el atajo que suelen tomar las plantillas típicas; col. C = lo '
             'que hace ESTE libro; col. D = por qué. Sin historia de versiones («v1.1», «antes», «ahora»), sin cifras '
             'del caso (cambian al recalcular: describe el método) y sin normativa española')
PISTA_D19 = ('D19: tabla «INDUSTRY BENCHMARKS» (Benchmark | Value | Source | Note). Valor = rango US de '
             'F1-research-us.md §3-§4 en formato US («28-35%», «$10-$15», sin € ni coma decimal); Source = el '
             'dominio de la fuente de research (p. ej. «Toast (pos.toasttab.com)») o «Kit estimate»; NUNCA «Fichero '
             'v1.1» ni «BOE». Filas de convenio/SMI → suelo de salario mínimo federal + estatal (DOL)')

# Pistas por celda (instrucciones de la SPEC que solo aplican a una celda)
CELDA_PISTAS = OrderedDict([
    (('FTP', '0. Supuestos', 'A2'), 'D8: añade una vez que los importes van «in your currency» (sin símbolo)'),
    (('CAFP', '0. Supuestos', 'A2'), 'D8: añade una vez que los importes van «in your currency» (sin símbolo)'),
    (('FTP', 'Instrucciones', 'A8'), 'SPEC §4.1 línea 4: sin divisores 1,10/1,21 → «divide a menu price by 1 + your '
                                     'sales tax rate»; P&L excl. sales tax, cash flow incl.'),
    (('CAFP', 'Instrucciones', 'A8'), 'SPEC §4.1 línea 4: sin divisores 1,10/1,21 → «divide a menu price by 1 + your '
                                      'sales tax rate»; P&L excl. sales tax, cash flow incl.'),
    (('FTP', 'Instrucciones', 'A12'), 'SPEC §4.1 línea 8 (I1): el Word del kit usa las MISMAS cifras de ejemplo que '
                                      'este libro; si cambias el libro, actualiza las cifras del Word'),
    (('CAFP', 'Instrucciones', 'A13'), 'SPEC §4.1 línea 8 (I1): el Word del kit usa las MISMAS cifras de ejemplo que '
                                       'este libro; si cambias el libro, actualiza las cifras del Word'),
    (('FTP', 'Tesorería 12 meses', 'O5'), 'SPEC §4.1: estacionalidad «your area» (sin agosto/costa de España)'),
    (('CAFP', 'Tesorería 12 meses', 'O5'), 'SPEC §4.1: estacionalidad «your area» (sin agosto/costa de España)'),
])

# ==========================================================================
# 8. Reglas US/UK con fuente (SPEC §3; research §3-§5, §7). clave → (valor, fuente). [kit estimate] = no es norma.
# ==========================================================================
REGLAS_US = OrderedDict([
    ('sales_tax', ('8% de ejemplo (state + local), sin crédito por compras; UK VAT 20%',
                   'Tax Foundation [8% = kit estimate]; HMRC VAT Notice 709/1')),
    ('input_tax', ('0 en US (el sales tax pagado es coste); UK 20% si está registrado', 'financial-kit D8')),
    ('payroll', ('cargas de empresa ≈ 10% (FICA 7.65% + FUTA/SUTA + workers\' comp)', 'IRS Pub. 15; Topic 759 [kit estimate]')),
    ('min_wage', ('$7.25/h federal → $15,080/año a 2,080 h; estatal/local suele ser mayor', 'dol.gov/agencies/whd/minimum-wage')),
    ('flsa', ('40 h semanales: umbral de horas extra', 'dol.gov/agencies/whd/flsa')),
    ('income_tax', ('25% efectivo de ejemplo (21% federal C corp + estatal); pass-through 0%', 'IRC §11(b) [25% = kit estimate]')),
    ('nol', ('pérdidas compensables hasta el 80% de la base (el modelo compensa el 100%, simplificación)', 'IRC §172')),
    ('loan', ('SBA 7(a): 10% nominal de ejemplo; plazo hasta 10 años equipo/circulante', '13 CFR 120.212-120.214 [10% = kit estimate]')),
    ('microloan', ('SBA Microloan hasta $50,000 vía intermediarios sin ánimo de lucro', 'sba.gov «Microloans»')),
    ('dscr', ('1.15× límite y 1.25× objetivo, sin «SBA minimum»', '[kit estimate]; práctica de banca')),
    ('equity', ('aportación propia habitual ≈ 10% en start-ups SBA; muchos piden 20-30%', 'SBA SOP 50 10 vigente [orientativo]')),
    ('vidas', ('lineal de libros: FT 7/7, CAF 10/7; MACRS / Section 179 solo en nota', 'financial-kit D10; IRS Pub. 946')),
    ('allergens_us', ('9 alérgenos mayores (sésamo desde 2023)', 'fda.gov food allergies; FASTER Act')),
    ('allergens_uk', ('14 alérgenos', 'food.gov.uk')),
    ('cdl', ('CDL solo si GVWR ≥ 26,001 lb', '49 CFR 383.5 / 383.91')),
    ('uk_registro', ('registrar el negocio de alimentos 28 días antes de empezar', 'gov.uk/food-business-registration')),
    ('uk_fhrs', ('Food Hygiene Rating 0-5 (Escocia: FHIS, Pass / Improvement Required)', 'food.gov.uk food hygiene rating scheme')),
    ('sqft', ('1 m² = 10.764 sq ft', 'NIST SP 811')),
    ('galon', ('1 US gal = 3.785 L', 'NIST SP 811')),
    ('tarjeta', ('3.5% tasa combinada (porcentaje + fijo por cobro, que pesa más en tickets pequeños); comprobar '
                 'el extracto del procesador', '[kit estimate] (revisión final M1)')),
    ('delivery', ('comisión de apps 15-30% (25% en el ejemplo)', '[kit estimate]')),
])

# Glosario §2.3 (se suma al del financial-kit, que manda)
GLOSARIO = OrderedDict([
    ('plan de negocio', 'business plan'), ('cliente/día', 'customers per day'),
    ('ticket medio sin IVA', 'average check (excl. sales tax)'), ('PVP con IVA', 'menu price incl. sales tax'),
    ('aparcamiento y base del vehículo', 'commissary & truck parking'), ('venta ambulante', 'mobile vending'),
    ('gestor/gestoría', 'accountant / bookkeeper'), ('autónomo', 'sole proprietor'), ('SL', 'LLC'),
    ('alta en Hacienda', "state tax registration / seller's permit"),
    ('Seguridad Social (empresa)', 'employer payroll taxes'), ('SMI', 'minimum wage'),
    ('convenio', 'state / local minimum wage'), ('carencia', 'interest-only period'),
    ('fondo de maniobra', 'working capital'), ('DDD', 'pest control'),
    ('gestor de residuos', 'licensed waste / grease hauler'), ('carnet de manipulador', 'food handler card'),
    ('APPCC', 'HACCP-based food safety plan'), ('registro sanitario', 'health permit'),
    ('Inversión Inicial', 'startup costs'), ('PyG / cuenta de resultados', 'P&L'),
    ('punto de equilibrio', 'break-even'), ('tesorería', 'cash flow'), ('Financiación', 'financing'),
    ('Personal (hoja)', 'Staffing'), ('plantilla', 'staff / headcount'), ('jornada', 'FTE (full-time equivalent)'),
    ('pagas', 'pay periods'), ('IVA (ventas)', 'sales tax (US) / VAT (UK)'),
    ('IVA soportado', 'recoverable input tax'), ('Impuesto de Sociedades', 'income tax'),
    ('bases negativas', 'net operating losses (NOL)'), ('préstamo bancario', 'bank loan / SBA loan'),
    ('recursos propios', "owner's equity"), ('m²', 'sq ft'), ('litros', 'gallons'),
])

# ==========================================================================
# 9. Docx (D22, D23). Tokens de cifras: los traductores escriben {{nombre}}; `ensamblar_docx.py` los rellena desde
#    `cifras_caso.json` (lo genera aplicar_en.py con `calcular_cifras()` sobre la caché del xlsx EN calibrado).
#    Fuente: ('fila', hoja ES, rótulo ES de la col. A (str o {plan: str}), columna) — extraer_textos.py la resuelve
#    a celda en cada plan (censo_es.json → tokens_docx) — o ('deriv', función de otras cifras).
# ==========================================================================
FMT = ('usd', 'usd2', 'int', 'ceil', 'pct0', 'pct1', 'x2', 'dec1', 'dec2', 'anos1')
_S, _I, _P, _E, _T, _F, _R, _ES, _N = ('0. Supuestos', 'Inversión Inicial', 'PyG 3 Años', 'Punto Equilibrio',
                                       'Tesorería 12 meses', 'Financiación', 'Personal', 'Escenarios', 'Instrucciones')
TOKENS_DOCX = OrderedDict([
    # actividad
    ('clientes_dia', ('int', ('fila', _S, 'Clientes/día (media del año 1)', 'B'), 'average customers per day, year 1')),
    ('clientes_a2', ('int', ('fila', _P, 'Clientes/día (media)', 'C'), 'average customers per day, year 2')),
    ('clientes_a3', ('int', ('fila', _P, 'Clientes/día (media)', 'D'), 'average customers per day, year 3')),
    ('ticket', ('usd2', ('fila', _S, 'Ticket medio SIN IVA (€)', 'B'), 'average check, excl. sales tax')),
    ('precio_menu', ('usd2', ('fila', _S, 'PVP equivalente con IVA (calculado)', 'B'), 'average menu price incl. sales tax')),
    ('dias', ('int', ('fila', _S, 'Días de apertura al año', 'B'), 'opening days per year')),
    ('crecimiento_a2', ('pct0', ('fila', _S, 'Crecimiento de volumen del año 2', 'B'), 'volume growth in year 2')),
    ('crecimiento_a3', ('pct0', ('fila', _S, 'Crecimiento de volumen del año 3', 'B'), 'volume growth in year 3')),
    ('sales_tax', ('pct0', ('fila', _S, 'IVA repercutido de restauración', 'B'), 'example sales tax rate')),
    ('impuesto_beneficios', ('pct0', ('fila', _S, 'Impuesto de Sociedades, tipo general', 'B'), 'effective income tax rate')),
    ('food_cost_comida', ('pct0', ('fila', _S, 'Coste de mercancía sobre las ventas de COMIDA', 'B'), 'food cost on food sales')),
    ('food_cost_bebida', ('pct0', ('fila', _S, 'Coste de mercancía sobre las ventas de BEBIDA', 'B'), 'beverage cost on beverage sales')),
    ('consumibles_pct', ('pct0', ('fila', _S, 'Consumibles sobre ventas', 'B'), 'packaging and disposables, % of sales')),
    ('tarjeta_pct', ('pct1', ('fila', _S, 'Comisión de los medios de pago', 'B'), 'card processing fees, % of sales')),
    ('cargas_pct', ('pct0', ('fila', _S, 'Seguridad Social a cargo de la empresa', 'B'), 'employer payroll taxes, % of wages')),
    ('salario_minimo_federal', ('usd', ('fila', _S, 'SMI anual de referencia (€)', 'B'), 'federal minimum wage, annual full-time')),
    ('alquiler_mes', ('usd', ('fila', _S, 'Alquiler mensual del local (€)', 'B'),
                      'rent per month (food truck: commissary + truck parking)')),
    ('suministros_mes', ('usd', ('fila', _S, 'Suministros mensuales de luz, agua y gas (€)', 'B'), 'utilities per month')),
    ('seguros_anual', ('usd', ('fila', _S, 'Seguros (€/año)', 'B'), 'insurance per year')),
    ('vida_obra', ('int', ('fila', _S, 'Vida útil de obra e instalaciones (años)', 'B'),
                   'useful life of build-out (food truck: truck and build-out), years')),
    ('vida_equipo', ('int', ('fila', _S, 'Vida útil de maquinaria y mobiliario (años)', 'B'), 'useful life of equipment, years')),
    # P&L
    ('ventas_a1', ('usd', ('fila', _P, 'INGRESOS TOTALES (sin IVA)', 'B'), 'revenue year 1, excl. sales tax')),
    ('ventas_a2', ('usd', ('fila', _P, 'INGRESOS TOTALES (sin IVA)', 'C'), 'revenue year 2')),
    ('ventas_a3', ('usd', ('fila', _P, 'INGRESOS TOTALES (sin IVA)', 'D'), 'revenue year 3')),
    ('ventas_dia', ('usd', ('deriv', lambda c: c['ventas_a1'] / c['dias']), 'average sales per opening day, year 1')),
    ('cogs_pct', ('pct1', ('fila', _P, 'Coste de mercancía / Ventas', 'B'), 'cost of goods (food + beverage), % of sales')),
    ('margen_bruto_pct', ('pct1', ('fila', _P, 'Margen bruto / Ventas', 'B'), 'gross margin, % of sales')),
    ('personal_a1', ('usd', ('fila', _P, 'Personal (nóminas + Seguridad Social)', 'B'), 'labor cost year 1 incl. payroll taxes')),
    ('personal_pct', ('pct1', ('fila', _P, 'Coste de personal / Ventas', 'B'), 'labor cost, % of sales')),
    ('ocupacion_anual', ('usd', ('fila', _P, {'ft': 'Aparcamiento y base del vehículo', 'caf': 'Alquiler del local'}, 'B'),
                         'occupancy per year (food truck: commissary + parking; coffee shop: rent)')),
    ('ocupacion_pct', ('pct1', ('fila', _P, {'ft': 'Aparcamiento / Ventas', 'caf': 'Alquiler / Ventas'}, 'B'),
                       'occupancy, % of sales')),
    ('marketing_anual', ('usd', ('fila', _P, 'Marketing y publicidad', 'B'), 'marketing budget per year')),
    ('costes_fijos_a1', ('usd', ('fila', _P, 'TOTAL COSTES FIJOS', 'B'), 'total fixed costs year 1 (incl. depreciation and interest)')),
    ('amortizacion_anual', ('usd', ('fila', _P, 'Amortización del inmovilizado', 'B'), 'annual depreciation')),
    ('intereses_a1', ('usd', ('fila', _P, 'Gastos financieros (intereses del préstamo)', 'B'), 'loan interest year 1')),
    ('resultado_antes_impuestos_a1', ('usd', ('fila', _P, 'RESULTADO ANTES DE IMPUESTOS', 'B'), 'pre-tax profit year 1')),
    ('ebitda_a1', ('usd', ('deriv', lambda c: c['resultado_antes_impuestos_a1'] + c['amortizacion_anual'] + c['intereses_a1']),
                   'EBITDA year 1 (pre-tax profit + depreciation + interest)')),
    ('ebitda_pct_a1', ('pct1', ('deriv', lambda c: (c['resultado_antes_impuestos_a1'] + c['amortizacion_anual'] +
                                                     c['intereses_a1']) / c['ventas_a1']), 'EBITDA margin year 1')),
    ('resultado_neto_a1', ('usd', ('fila', _P, 'RESULTADO NETO', 'B'), 'net profit year 1')),
    ('resultado_neto_a2', ('usd', ('fila', _P, 'RESULTADO NETO', 'C'), 'net profit year 2')),
    ('resultado_neto_a3', ('usd', ('fila', _P, 'RESULTADO NETO', 'D'), 'net profit year 3')),
    ('margen_neto_a1', ('pct1', ('fila', _P, 'Resultado neto / Ventas', 'B'), 'net margin year 1')),
    # inversión y financiación
    ('inversion_total', ('usd', ('fila', _I, 'INVERSIÓN TOTAL (suma de los bloques)', 'B'), 'total startup investment incl. working capital')),
    ('capex', ('usd', ('fila', _I, 'CAPEX (inversión sin el fondo de maniobra)', 'B'), 'capex (startup costs without working capital)')),
    ('fondo_maniobra', ('usd', ('fila', _I, 'Colchón operativo hasta alcanzar el equilibrio', 'B'), 'working capital reserve')),
    ('necesidad_caja', ('usd', ('fila', _I, 'NECESIDAD TOTAL DE CAJA AL ARRANQUE', 'B'), 'total cash needed at opening')),
    ('fondos_propios', ('usd', ('fila', _F, 'Recursos propios de los socios', 'B'), "owner's equity injection")),
    ('prestamo', ('usd', ('fila', _F, 'Importe del principal', 'B'), 'loan amount')),
    ('aportacion_pct', ('pct0', ('deriv', lambda c: c['fondos_propios'] / c['necesidad_caja']),
                        "owner's equity as % of the total cash needed")),
    ('tipo_prestamo', ('pct1', ('fila', _F, 'Tipo de interés nominal anual', 'B'), 'loan interest rate (nominal, not APR)')),
    ('plazo_prestamo', ('int', ('fila', _F, 'Plazo total (años)', 'B'), 'loan term, years')),
    ('cuota_anual', ('usd', ('fila', _F, 'Cuota anual durante la amortización', 'B'), 'annual loan payment')),
    ('dscr_a1', ('x2', ('fila', _F, 'DSCR del año 1', 'B'), 'DSCR year 1')),
    ('dscr_min', ('x2', ('fila', _F, 'DSCR mínimo de todo el cuadro', 'B'), 'lowest DSCR over the loan term')),
    ('dscr_limite', ('x2', ('fila', _S, 'DSCR mínimo aceptable', 'B'), 'DSCR limit used by the workbook')),
    ('dscr_objetivo', ('x2', ('fila', _S, 'DSCR objetivo (verde)', 'B'), 'DSCR target (typical lender guideline)')),
    # equilibrio y caja
    ('equilibrio_clientes_dia', ('ceil', ('fila', _E, 'Clientes necesarios al día', 'B'), 'break-even customers per day (accounting)')),
    ('equilibrio_ventas', ('usd', ('fila', _E, 'Ingresos necesarios al año', 'B'), 'break-even revenue per year (accounting)')),
    ('equilibrio_caja_clientes_dia', ('ceil', ('fila', _E, 'Clientes necesarios al día (caja)', 'B'),
                                      'cash break-even customers per day (loan payments in, depreciation out)')),
    ('equilibrio_caja_ventas', ('usd', ('fila', _E, 'Ingresos necesarios al año (caja)', 'B'), 'cash break-even revenue per year')),
    ('holgura_caja_pct', ('pct0', ('fila', _E, 'Holgura sobre el equilibrio de CAJA (%)', 'B'), 'margin of safety over cash break-even')),
    ('saldo_minimo', ('usd', ('fila', _T, 'Saldo mínimo del año', 'B'), 'lowest cash balance in year 1')),
    ('mes_saldo_minimo', ('int', ('fila', _T, 'Mes en el que la caja toca fondo', 'B'), 'month of the lowest cash balance')),
    ('payback', ('anos1', ('fila', _T, 'Payback del proyecto (años), antes de la deuda', 'B'),
                 'project payback before debt (renders "2.8 years" or "more than 3 years")')),
    ('payback_capex', ('anos1', ('fila', _T, 'Payback sobre el CAPEX, sin fondo de maniobra (años)', 'B'), 'payback on capex only')),
    # personal
    ('plantilla_puestos', ('int', ('fila', _R, 'TOTAL PLANTILLA', 'B'), 'number of people on the payroll')),
    ('plantilla_jornadas', ('dec2', ('fila', _R, 'TOTAL PLANTILLA', 'C'), 'staff in full-time equivalents (FTE)')),
    ('cobertura_horas', ('pct0', ('fila', _R, 'Cobertura (contratadas / necesarias)', 'B'), 'hours coverage (scheduled / needed)')),
    ('horas_semana', ('int', ('fila', _R, 'Jornada completa del convenio (horas/semana)', 'B'), 'full-time hours per week')),
    # escenarios
    ('pesimista_clientes', ('int', ('fila', _ES, 'Clientes/día', 'B'), 'pessimistic scenario: customers per day')),
    ('optimista_clientes', ('int', ('fila', _ES, 'Clientes/día', 'D'), 'optimistic scenario: customers per day')),
    ('pesimista_ventas', ('usd', ('fila', _ES, 'INGRESOS ANUALES (sin IVA)', 'B'), 'pessimistic scenario: annual revenue')),
    ('optimista_ventas', ('usd', ('fila', _ES, 'INGRESOS ANUALES (sin IVA)', 'D'), 'optimistic scenario: annual revenue')),
    ('pesimista_resultado', ('usd', ('fila', _ES, 'RESULTADO ANTES DE IMPUESTOS', 'B'), 'pessimistic scenario: pre-tax profit (negative = loss)')),
    ('optimista_resultado', ('usd', ('fila', _ES, 'RESULTADO ANTES DE IMPUESTOS', 'D'), 'optimistic scenario: pre-tax profit')),
    # ronda de arreglos (revisión final m1/m2): ticket y días de cada escenario (el lector reproduce los ingresos) y el
    # saldo de caja estimado al cierre del año 1 (en el FT pesimista se agota)
    ('pesimista_ticket', ('usd2', ('fila', _ES, 'Ticket medio sin IVA', 'B'), 'pessimistic scenario: average check excl. sales tax')),
    ('optimista_ticket', ('usd2', ('fila', _ES, 'Ticket medio sin IVA', 'D'), 'optimistic scenario: average check excl. sales tax')),
    ('pesimista_dias', ('int', ('fila', _ES, 'Días de apertura al año', 'B'), 'pessimistic scenario: operating days')),
    ('optimista_dias', ('int', ('fila', _ES, 'Días de apertura al año', 'D'), 'optimistic scenario: operating days')),
    ('pesimista_caja', ('usd', ('fila', _ES, 'Saldo de caja al cierre del año 1 (estimado, el mismo método en los tres '
                                'escenarios)', 'B'), 'pessimistic scenario: estimated year-1 closing cash (negative = cash runs out)')),
    # partidas de inversión
    ('inv_lanzamiento', ('usd', ('fila', _I, {'ft': 'Marketing lanzamiento', 'caf': 'Marketing lanzamiento (web + RRSS + Google)'}, 'B'),
                         'launch marketing budget')),
    ('inv_stock', ('usd', ('fila', _I, {'ft': 'Stock inicial de producto', 'caf': 'Stock inicial (café, leche, bollería, etc.)'}, 'B'),
                   'opening inventory')),
    ('inv_vehiculo', ('usd', ('fila', _I, {'ft': 'Vehículo food truck (nuevo o segunda mano)'}, 'B'), 'used food truck')),
    ('inv_adaptacion', ('usd', ('fila', _I, {'ft': 'Adaptación y rotulación vehículo'}, 'B'), 'truck build-out and wrap')),
    ('inv_equipo_cocina', ('usd', ('fila', _I, {'ft': 'Equipamiento cocina móvil'}, 'B'), 'kitchen and refrigeration equipment')),
    ('inv_generador', ('usd', ('fila', _I, {'ft': 'Generador eléctrico o conexión'}, 'B'), 'generator')),
    ('inv_permisos', ('usd', ('fila', _I, {'ft': 'Licencias y permisos (promedio)'}, 'B'), 'first-year permits and licenses')),
    ('inv_obra', ('usd', ('fila', _I, {'caf': 'Obra civil y adecuación local'}, 'B'), 'construction / build-out')),
    ('inv_instalaciones', ('usd', ('fila', _I, {'caf': 'Instalación eléctrica + climatización'}, 'B'), 'electrical + HVAC')),
    ('inv_espresso', ('usd', ('fila', _I, {'caf': 'Máquina de café espresso profesional'}, 'B'), 'two-group espresso machine')),
    ('inv_mobiliario', ('usd', ('fila', _I, {'caf': 'Mobiliario sala (mesas + sillas)'}, 'B'), 'dining room furniture')),
    ('inv_barra', ('usd', ('fila', _I, {'caf': 'Barra y mostrador expositor'}, 'B'), 'counter and pastry display bar')),
    ('inv_terraza', ('usd', ('fila', _I, {'caf': 'Terraza (mobiliario + toldo)'}, 'B'), 'patio / sidewalk seating')),
    ('aforo', ('int', ('fila', _S, {'caf': 'Aforo del local (plazas sentadas + barra)'}, 'B'), 'seats (tables + bar)')),
    ('rotaciones_dia', ('dec1', ('fila', _S, {'caf': 'Rotaciones al día implícitas (calculado)'}, 'B'), 'implied seat turns per day')),
])


def tokens_de(plan):
    """Tokens válidos para un plan (los de rótulo {plan: …} solo existen donde tienen rótulo)."""
    out = OrderedDict()
    for t, (fmt, src, desc) in TOKENS_DOCX.items():
        if src[0] == 'fila' and isinstance(src[2], dict) and plan not in src[2]:
            continue
        out[t] = (fmt, src, desc)
    return out


RX_TOKEN = re.compile(r'\{\{([a-z0-9_]+)\}\}')
_PAYBACK_TXT = {'Over 3 years': 'more than 3 years', 'Más de 3 años': 'more than 3 years'}


def formatear(valor, fmt):
    """Cifra → texto US para el docx. Importes con «$» (texto de un plan para un lector de EE. UU., SPEC §5)."""
    if isinstance(valor, str):
        if fmt == 'anos1':
            return _PAYBACK_TXT.get(valor, valor)
        return valor
    if valor is None:
        raise ValueError('cifra vacía')
    v = float(valor)
    neg = v < 0
    a = abs(v)
    if fmt == 'usd':
        s = '${:,.0f}'.format(a)
    elif fmt == 'usd2':
        s = '${:,.2f}'.format(a)
    elif fmt == 'int':
        s = '{:,.0f}'.format(a)
    elif fmt == 'ceil':
        s = '{:,.0f}'.format(math.ceil(a - 1e-9))
    elif fmt == 'pct0':
        s = '{:.0f}%'.format(a * 100)
    elif fmt == 'pct1':
        s = '{:.1f}%'.format(a * 100)
    elif fmt == 'x2':
        s = '{:.2f}×'.format(a)
    elif fmt == 'dec1':
        s = '{:,.1f}'.format(a)
    elif fmt == 'dec2':
        s = '{:,.2f}'.format(a)
    elif fmt == 'anos1':
        s = '{:.1f} years'.format(a)
    else:
        raise ValueError('formato desconocido: ' + fmt)
    return ('-' + s) if neg else s


def calcular_cifras(plan, leer, coords):
    """Cifras del caso para `cifras_caso.json`. leer(hoja_es, celda) → valor de la caché del xlsx EN (quien llama
    traduce la hoja); coords = censo_es.json['tokens_docx'][plan]. Devuelve {token: {valor, fmt, texto, hoja_es,
    celda}} con las derivadas calculadas al final."""
    base, out = {}, OrderedDict()
    toks = tokens_de(plan)
    for t, (fmt, src, desc) in toks.items():
        if src[0] != 'fila':
            continue
        hoja, celda = coords[t]['hoja_es'], coords[t]['celda']
        v = leer(hoja, celda)
        base[t] = v
        out[t] = OrderedDict([('valor', v), ('fmt', fmt), ('texto', formatear(v, fmt)), ('hoja_es', hoja),
                              ('celda', celda)])
    for t, (fmt, src, desc) in toks.items():
        if src[0] == 'deriv':
            v = src[1](base)
            out[t] = OrderedDict([('valor', v), ('fmt', fmt), ('texto', formatear(v, fmt)), ('deriv', True)])
    return out


# Encabezados D22 (Heading 1) e índice: «N. Título» / «N.  Título»
ENCABEZADOS_EN = ['Executive Summary', 'Concept & Value Proposition', 'Market Analysis', 'Competitive Analysis',
                  'Marketing Plan', 'Operations Plan', 'Management & Staffing', 'Financial Plan',
                  'Legal Requirements, Permits & Licenses', 'Conclusions & Action Plan']
AVISO_DOCX_EN = ('This document was prepared by John Guerrero for AI Chef Pro for planning and information purposes. '
                 'Its figures come from the example case in the Excel workbook of this kit and from the public sources '
                 'named in the text; your results will differ. It is not financial, legal or tax advice: requirements '
                 'vary by state, county and city — check with your accountant and local authorities.')
PRODUCTOS_CIERRE = ['Food Cost Kit Pro', 'HACCP Food Safety Kit Pro', 'Restaurant Staff Scheduling Kit Pro',
                    'Gastro Pro Prompts eBook (300+ AI prompts)']
# Párrafos que escribe `ensamblar_docx.py` (portada, aviso, índice, encabezados, cierre — D22, D23). Índices = ES.
DOCX_FIJOS = OrderedDict([
    ('ft', OrderedDict([
        (6, 'BUSINESS PLAN'), (7, 'Food Truck'),
        (9, 'Complete financial plan, market analysis,\noperations strategy and startup checklist\nto launch a food '
            'truck in the United States'),
        (14, 'AI Chef Pro'), (15, 'aichef.pro'), (16, 'US edition with UK notes · 2026'),
        (18, 'LEGAL NOTICE'), (19, AVISO_DOCX_EN), (20, '© 2026 AI Chef Pro — aichef.pro'), (22, 'CONTENTS'),
        (35, 'Appendices: Food Truck Financial Projections (Excel) + Food Truck Startup Checklist (Excel)'),
        (123, 'AI Chef Pro'), (124, 'Templates and AI tools for food businesses'), (126, 'aichef.pro\ninfo@aichef.pro'),
        (127, '→ %s · %s' % tuple(PRODUCTOS_CIERRE[:2])), (128, '→ ' + PRODUCTOS_CIERRE[2]),
        (129, '→ ' + PRODUCTOS_CIERRE[3]), (130, '\n' + URL_TIENDA),
    ])),
    ('caf', OrderedDict([
        (6, 'BUSINESS PLAN'), (7, 'Coffee Shop'),
        (9, 'Complete financial plan, market analysis,\noperations strategy and opening checklist\nto open a coffee '
            'shop with brunch in the United States'),
        (14, 'AI Chef Pro'), (15, 'aichef.pro'), (16, 'US edition with UK notes · 2026'),
        (18, 'LEGAL NOTICE'), (19, AVISO_DOCX_EN), (20, '© 2026 AI Chef Pro — aichef.pro'), (22, 'CONTENTS'),
        (35, 'Appendices: Coffee Shop Financial Projections (Excel) + Coffee Shop Opening Checklist (Excel)'),
        (156, 'AI Chef Pro'), (157, 'Templates and AI tools for food businesses'),
        (159, 'Built on more than 29 years in professional kitchens and restaurant consulting.\nAI agents, templates '
              'and digital products for restaurants, cafes, bars and food trucks.\n\naichef.pro\ninfo@aichef.pro'),
        (161, 'Complete this plan with our digital products:'),
        (162, '→ ' + PRODUCTOS_CIERRE[0]), (163, '→ ' + PRODUCTOS_CIERRE[1]), (164, '→ ' + PRODUCTOS_CIERRE[2]),
        (165, '→ ' + PRODUCTOS_CIERRE[3]), (166, '\n' + URL_TIENDA),
    ])),
])
# Índices de los 10 encabezados y de las 10 líneas del índice (se verifican contra el docx ES en extraer_docx.py)
DOCX_ENCABEZADOS = OrderedDict([('ft', [37, 44, 51, 59, 67, 76, 83, 91, 100, 109]),
                                ('caf', [37, 47, 56, 74, 82, 94, 106, 132, 135, 143])])
DOCX_INDICE = OrderedDict([('ft', list(range(24, 34))), ('caf', list(range(24, 34)))])
# CAF (defecto I7 del ES): §8 dice «[Contenido pendiente de generacion]» y el bloque financiero vive dentro de §7
# (párrafos 110-119). La EN lo MUEVE a §8 sin cambiar el nº de párrafos ni la secuencia de estilos (todos Normal):
# tras el 133 (separador de §8) van 110 (subtítulo), 134 (resumen nuevo con tokens) y 111-119.
DOCX_MOVER = OrderedDict([('caf', OrderedDict([('despues_de', 133), ('orden', [110, 134] + list(range(111, 120)))]))])
# CAF 134 lleva el run en ROJO (cc0000: era un aviso de «pendiente»): el ensamblador le copia el formato del run de un
# párrafo de cuerpo normal (excepción de formato declarada; G8 la admite por nombre)
DOCX_RPR_DE = OrderedDict([('caf', OrderedDict([(134, 112)]))])
# Notas por párrafo para los traductores (BRIEF-TRADUCTORES.md las repite; extraer_docx.py las copia a docx_es_*.json)
DOCX_NOTAS = OrderedDict([
    ('ft', OrderedDict([
        (102, 'D22 §9 párrafo 1/7: el marco US por niveles (federal, estatal, condado, ciudad): qué regula cada nivel y '
              'que un food truck tramita en CADA ciudad donde vende. Fuente: sba.gov (apply for licenses and permits)'),
        (103, '§9 2/7: mobile vending license de la ciudad (por ciudad, intransferible, una por unidad); zonas y reglas '
              'de aparcamiento; permiso escrito para vender en propiedad privada. Ejemplo citable: chicago.gov'),
        (104, '§9 3/7: health permit de unidad móvil + plan review del health department + commissary agreement (la '
              'mayoría de jurisdicciones lo exige) + inspección. Ejemplo citable: cabq.gov (Albuquerque)'),
        (105, '§9 4/7: higiene y alérgenos: Certified Food Protection Manager (p. ej. ServSafe) + food handler cards '
              'según estado; FDA Food Code adoptado por cada estado; 9 alérgenos mayores (sésamo desde 2023). Sustituye '
              'el párrafo del RGSEAA (que en el ES además era erróneo)'),
        (106, "§9 5/7: seguros: general liability (a menudo $1M por siniestro, lo piden eventos y propietarios) + "
              "commercial auto + workers' comp; el ejemplo del libro presupuesta {{seguros_anual}} al año"),
        (107, '§9 6/7: vehículo: registro y título en el DMV del estado; CDL SOLO si GVWR ≥ 26,001 lb (49 CFR 383); '
              'inspección del fire marshal (propano, extintores, supresión; NFPA 96 donde se adopte)'),
        (108, '§9 7/7: varias ciudades y eventos (cada ciudad su licencia; temporary food event permits; los '
              'organizadores piden seguro y permisos) y, como ÚLTIMA frase o frases del párrafo, la nota UK: registrar '
              'el negocio de alimentos en el council 28 días antes de empezar (gratis), street trading licence/consent '
              'de cada council, Food Hygiene Rating 0-5, VAT 20 % en comida caliente y en el local, 14 alérgenos'),
    ])),
    ('caf', OrderedDict([
        (110, 'Este párrafo y el bloque 111-119 PASAN a la sección 8 (el ensamblador los mueve: en el ES el bloque '
              'financiero estaba dentro de §7 y §8 decía «[Contenido pendiente de generacion]»). 110 = subtítulo '
              'corto de §8, p. ej. «Financial plan at a glance»'),
        (134, 'En el ES: «[Contenido pendiente de generacion]». En la EN va JUSTO DESPUÉS del subtítulo 110: escribe el '
              'resumen del plan financiero (inversión, fondos propios + préstamo, ventas año 1-3, EBITDA y resultado, '
              'equilibrio, DSCR, saldo mínimo de caja) SOLO con tokens'),
        (111, 'Subtítulo (pasa a §8): «Startup costs and funding» o similar'),
        (114, 'Subtítulo (pasa a §8): «Sales forecast and margins» o similar'),
        (117, 'Subtítulo (pasa a §8): «Scenario analysis» o similar'),
        (108, 'Subtítulo de §7 (texto plano, sin negrita: es un párrafo Normal de una línea)'),
        (120, 'Subtítulo de §7'), (123, 'Subtítulo de §7'), (126, 'Subtítulo de §7'), (129, 'Subtítulo de §7'),
        (137, 'D22 §9 1/6: estructura (LLC vs sole proprietorship; Secretary of State; registered agent; operating '
              'agreement) + EIN gratis en el IRS. Fuente: sba.gov (choose a business structure), irs.gov'),
        (138, "§9 2/6: registros de impuestos y licencias: seller's permit / state sales tax registration, business "
              'license de la ciudad o el condado, DBA si aplica (sustituye el IAE)'),
        (139, '§9 3/6: zoning check ANTES de firmar el alquiler + building permits de la obra + certificate of '
              'occupancy (sustituye la licencia de actividad)'),
        (140, '§9 4/6: health permit del food establishment con plan review e inspección + Certified Food Protection '
              'Manager + food handler cards + plan de seguridad alimentaria basado en HACCP donde se exija + 9 alérgenos '
              'mayores (sustituye el APPCC). Fuente: fda.gov (Food Code)'),
        (141, '§9 5/6: sidewalk café permit (terraza) + sign permit + liquor license SOLO si sirves alcohol (pídelo '
              'pronto: puede tardar meses)'),
        (142, "§9 6/6: empleo y seguros: OSHA basics + workers' comp + I-9 / W-4 / new-hire reporting + general "
              'liability y BOP; como ÚLTIMA frase o frases, la nota UK: registrar el negocio de alimentos en el council '
              '28 días antes de abrir (gratis), Food Hygiene Rating 0-5, VAT 20 % en el local y en comida caliente, 14 '
              'alérgenos, permiso del council para mesas en la calle'),
    ])),
])
# Tandas de traducción del docx (BRIEF-TRADUCTORES.md): secciones por tanda
DOCX_TANDAS = OrderedDict([('ft_a', ('ft', [1, 2, 3, 4, 5])), ('ft_b', ('ft', [6, 7, 8, 9, 10])),
                           ('caf_a', ('caf', [1, 2, 3, 4, 5])), ('caf_b', ('caf', [6, 7, 8, 9, 10]))])


def textos_fijos_docx(plan):
    """Índice ES → texto EN de todo lo que escribe el ensamblador (DOCX_FIJOS + encabezados + índice)."""
    d = OrderedDict(DOCX_FIJOS[plan])
    for n, (ih, ii) in enumerate(zip(DOCX_ENCABEZADOS[plan], DOCX_INDICE[plan]), 1):
        d[ih] = '%d. %s' % (n, ENCABEZADOS_EN[n - 1])
        d[ii] = '%d.  %s' % (n, ENCABEZADOS_EN[n - 1])
    return d


# ==========================================================================
# 10. Detectores (compartidos por check_textos.py y ensamblar_docx.py; copia del de financial-kit/gates_en.py)
# ==========================================================================
SIGNOS_ES = set('¿¡«»€ºª')
RX_TILDE = re.compile(r'[áéíóúüñÁÉÍÓÚÜÑ]')
LISTA_BLANCA_TILDE = {'café', 'cafés', 'entrée', 'entrées', 'sauté', 'sautéed', 'purée', 'jalapeño', 'jalapeños',
                      'crème', 'brûlée', 'résumé', 'résumés', 'naïve', 'açaí', 'piñata', 'señor', 'jalapeños'}
FUNCIONALES_ES = {
    'de', 'del', 'las', 'el', 'los', 'que', 'con', 'para', 'por', 'una', 'unos', 'unas', 'es', 'al', 'se',
    'su', 'sus', 'como', 'pero', 'muy', 'cuando', 'donde', 'porque', 'este', 'esta', 'esto', 'estos', 'desde',
    'hasta', 'sobre', 'entre', 'tu', 'tus', 'mis', 'hoja', 'pestaña', 'celda', 'celdas', 'sin', 'ejemplo',
    'cada', 'debe', 'hay', 'puede', 'lo', 'les', 'todo', 'toda', 'todos', 'nunca', 'siempre', 'también', 'año',
    'años', 'mes', 'meses', 'según',
}
OFICIO_ES = {
    'ventas', 'venta', 'facturacion', 'gastos', 'gasto', 'ingresos', 'ingreso', 'coste', 'costes', 'prestamo',
    'inversion', 'tesoreria', 'amortizacion', 'beneficio', 'beneficios', 'impuesto', 'impuestos', 'cuota', 'saldo',
    'cobros', 'cobro', 'pagos', 'proveedores', 'proveedor', 'alquiler', 'suministros', 'nominas', 'nomina',
    'plantilla', 'plantillas', 'cubiertos', 'comedor', 'barra', 'anual', 'mensual', 'resumen', 'datos', 'escenarios',
    'escenario', 'supuestos', 'flujo', 'obra', 'licencias', 'licencia', 'equipamiento', 'mobiliario', 'fondos',
    'propios', 'deuda', 'banco', 'garantias', 'pendiente', 'presupuesto', 'partida', 'concepto', 'tasa', 'plazo',
    'cuadro', 'intereses', 'objetivo', 'limite', 'estado', 'fase', 'tarea', 'tareas', 'responsable', 'notas',
    'gestoria', 'gestor', 'socio', 'socios', 'promotor', 'aforo', 'plazas', 'superficie', 'ubicacion', 'apertura',
    'constitucion', 'seguro', 'seguros', 'contrato', 'notario', 'registro', 'liquidacion', 'tarjeta', 'compras',
    'bebida', 'comida', 'otros', 'umbral', 'minimo', 'semaforo', 'dias', 'horas', 'margen', 'punto', 'equilibrio',
    'financiacion', 'cocina', 'hosteleria', 'negocio', 'materia', 'carencia', 'traspaso', 'fianza', 'aval',
    'ayuntamiento', 'vehiculo', 'furgoneta', 'cafeteria', 'semana', 'semanas', 'titular', 'aseguradora', 'trámite',
    'tramite', 'personas', 'jornada', 'pagas', 'clientes', 'cliente',
}
RX_URL = re.compile(r'https?://\S+|\b[\w-]+(?:\.[\w-]+)*\.(?:es|com|pro|gov|org|uk|co|net|io)\b[/\w.#-]*', re.I)
RX_PALABRA = re.compile(r"[A-Za-zÀ-ÖØ-öø-ÿ]+")
# Topónimos de EE. UU. con artículos españoles («Los Angeles», «El Paso»…): no son restos de español
RX_LUGARES_US = re.compile(r'\b(?:Los Angeles|Las Vegas|Las Cruces|El Paso|El Segundo|Los Gatos|La Jolla|Del Mar|'
                           r'San (?:Diego|Francisco|Antonio|Jose|Juan)|Santa (?:Fe|Monica|Barbara|Clara|Cruz))\b')
RX_SIGLAS_ES = re.compile(r'\b(?:TPV|IVA|CIF|NIF|IRPF|TIN|TAE|SGR|ICO|ENISA|RETA|CCC|BAI|LIVA|LIS|IAE|SMI|CCAA|RGSEAA|'
                          r'ITV|LOPDGDD|RGPD|AEPD|ROESB|SGAE|AGEDI|FNMT|OEPM|LAU|APPCC|DDD|PRL|BOE)\b')
RX_NORMA_ES = re.compile(r'Veri\*factu|[Mm]odelo 303|\bS\.?L\.?\b(?=\s|$|[,.;)])|[Aa]utónom|[Cc]onvenio|'
                         r'Seguridad Social|Hacienda|Registro Mercantil|Comunidad Autónoma|Estatuto de los Trabajadores')
RX_COMA_DECIMAL = re.compile(r'(?<![\d,])\d+,\d{1,2}(?![\d])')
RX_MONEDA_ES = re.compile(r'€|\bEUR\b|\beuros?\b', re.I)
RX_MARCA_ES = re.compile(r'ChefBusiness|chefbusiness\.co', re.I)
RX_ID_INTERNO = re.compile(r'\[(?:fuente|estimado|derivado|kit estimate|kit benchmark)[^\]]*\]|\bSPEC\b|'
                           r'\b(?:D|E|I)\d{1,2}\b(?=[:,)])|\bF1-research|\bcifras_caso\b', re.I)


def restos_espanol(t):
    """Lista de (motivo, contexto) con restos de español en un texto EN (vacía = limpio)."""
    out = []
    limpio = RX_LUGARES_US.sub(' ', RX_URL.sub(' ', t))
    for ch in limpio:
        if ch in SIGNOS_ES:
            out.append(('signo «%s»' % ch, t[:80]))
    for m in RX_PALABRA.finditer(limpio):
        w = m.group(0)
        wl = w.lower()
        ctx = limpio[max(0, m.start() - 30):m.end() + 30].replace('\n', ' ')
        sin = unicodedata.normalize('NFKD', wl).encode('ascii', 'ignore').decode()
        if RX_TILDE.search(w) and wl not in LISTA_BLANCA_TILDE:
            out.append(('tilde/eñe «%s»' % w, ctx))
        elif wl in FUNCIONALES_ES and w not in ('NO', 'SE', 'SI', 'S', 'TU', 'ES', 'AL', 'DE', 'EL', 'LOS'):
            out.append(('funcional ES «%s»' % w, ctx))
        elif sin in OFICIO_ES and wl not in LISTA_BLANCA_TILDE:
            out.append(('oficio ES «%s»' % w, ctx))
    for rx, nombre in ((RX_COMA_DECIMAL, 'coma decimal'), (RX_SIGLAS_ES, 'sigla ES'), (RX_NORMA_ES, 'normativa ES'),
                       (RX_MONEDA_ES, 'moneda ES'), (RX_MARCA_ES, 'marca ES')):
        m = rx.search(t)
        if m:
            out.append(('%s «%s»' % (nombre, m.group(0)), t[max(0, m.start() - 30):m.end() + 30]))
    if 'm²' in t:
        out.append(('m² (D14: sq ft)', t[:80]))
    return out


_RANGOS_NO_LATINOS = [(0x0370, 0x03FF), (0x0400, 0x052F), (0x0590, 0x05FF), (0x0600, 0x06FF), (0x0E00, 0x0E7F),
                      (0x1100, 0x11FF), (0x3040, 0x30FF), (0x3400, 0x4DBF), (0x4E00, 0x9FFF), (0xAC00, 0xD7AF),
                      (0xF900, 0xFAFF), (0xFF00, 0xFFEF)]


def no_latinos(t):
    return sorted({ch for ch in t if any(a <= ord(ch) <= b for a, b in _RANGOS_NO_LATINOS)})


# ==========================================================================
# 11. Fórmulas con texto (D7, D11) — ver `_formulas_en()` en formulas_en.py (se separan por tamaño)
# ==========================================================================
try:
    from formulas_en import FORMULAS_EN as _FX                                     # noqa: E402
    FORMULAS_EN.update(_FX)
except ImportError:                                                                # pragma: no cover
    pass


# ==========================================================================
# 12. Parámetros por conjunto que antes vivían escritos en los scripts (F2 tanda 1 de Restaurant + Bakery)
# ==========================================================================
FUENTES_CIFRAS = [os.path.join(AQUI, 'F1-research-us.md'), os.path.join(AQUI, 'SPEC.md')]   # cifras citables sin token
NOTAS_F2 = os.path.join(AQUI, 'F2-NOTAS.md')               # G7: cada calibración D15 citada aquí
DOCX_N_PARRAFOS = OrderedDict([('ft', 131), ('caf', 167)])  # SPEC D22
DOCX_ULTIMO_9 = OrderedDict([('ft', '108'), ('caf', '142')])  # último párrafo de §9: cierra con la nota UK
DOCX_RESUMEN_TOKENS = OrderedDict([('caf', ('134', 3))])    # párrafo que debe citar ≥ n tokens
PRESTAMO_SEMILLA = OrderedDict([('ft', 80000), ('caf', 140000)])
REUSO = OrderedDict()                                      # ES → EN ya traducido en otro conjunto (solo restpan)
CON_MARCA = None                                           # libros con la línea «More templates…» (None = todos)
# Autotest de gates_en.py: papeles de los libros (P1/P2 planes, C1/C2 checklists) y celdas que muta
AUTOTEST = OrderedDict([
    ('P1', 'FTP'), ('P2', 'CAFP'), ('C1', 'FTC'), ('C2', 'CAFC'),
    ('C1_hoja_papel', 'F2 - Vehículo'), ('C1_titulo', 'Food Truck Startup Checklist (68 Tasks)'),
    ('C2_hoja_no_latino', 'F5 - Marketing'), ('P1_celda_texto', 'A3'), ('P1_celda_formula', 'B9'),
    ('P1_merge', 'A30:C30'),
])


# ==========================================================================
# autotest y cruce con el censo
# ==========================================================================
RX_ES_CLAVE = re.compile(r'[áíóúñÁÍÓÚÑ¿¡€]|\b(?:de|del|los|las|con|sin|para|por|IVA)\b')


def necesita_clave(lit):
    return any(ch.isalpha() for ch in lit) or '€' in lit


def autotest():
    err = []
    for c, d in HOJAS_POR_LIBRO.items():
        ens = list(d.values())
        if len(set(ens)) != len(ens):
            err.append('pestañas EN duplicadas en ' + c)
        for en in ens:
            if len(en) > 31 or set(en) & HOJA_PROHIBIDOS or en.startswith("'"):
                err.append('pestaña inválida: ' + en)
        if list(d) != HOJAS_ES_POR_LIBRO[c] and sorted(d) != sorted(HOJAS_ES_POR_LIBRO[c]):
            err.append('HOJAS_POR_LIBRO y HOJAS_ES_POR_LIBRO no cuadran en ' + c)
    if len(set(FICHEROS.values())) != 6:
        err.append('FICHEROS debe tener 6 nombres EN distintos')
    for k, v in CLAVES.items():
        if '"' in v or RX_ES_CLAVE.search(v):
            err.append('comilla o español en una clave EN: %r' % v)
    for nombre, d in (('FIJOS', FIJOS), ('POR_CELDA', POR_CELDA)):
        for k, v in d.items():
            r = [x for x in restos_espanol(v) if not x[0].startswith('oficio')]
            if r or re.search(r'\d,\d{1,2}(?!\d)', v):
                err.append('%s con español o coma decimal: %r %s' % (nombre, v[:60], r[:1]))
            if RX_ID_INTERNO.search(v):
                err.append('%s con un ID interno: %r' % (nombre, v[:60]))
    for (c, h, x), (t, _e) in TEXTOS_NUEVOS.items():
        if restos_espanol(t) or RX_ID_INTERNO.search(t):
            err.append('TEXTOS_NUEVOS con español o ID interno: %r' % t[:60])
    for k, v in FORMULAS_EN.items():
        lits = re.findall(r'"((?:[^"]|"")*)"', v)
        for lit in lits:
            r = [x for x in restos_espanol(lit) if not x[0].startswith('oficio')]
            if r or '€' in lit:
                err.append('FORMULAS_EN %r con español: %r %s' % (k, lit[:50], r[:1]))
        for hes in todas_hojas_es():
            if hes not in HOJAS_POR_LIBRO[k[0]].values() and ("'%s'!" % hes in v or '%s!' % hes in v):
                err.append('FORMULAS_EN %r cita la pestaña ES %r' % (k, hes))
        if len(v) > 8192:
            err.append('FORMULAS_EN %r > 8192 caracteres' % (k,))
    for k, (viejo, nuevo) in PARCHES_FORMULA.items():
        if viejo == nuevo:
            err.append('PARCHE vacío: %r' % (k,))
    for k, (v, rot) in VALORES_EN.items():
        if not (v is None or v == CALIBRAR or isinstance(v, (int, float))):
            err.append('VALORES_EN no numérico: %r' % (k,))
    for k, (val, fuente) in REGLAS_US.items():
        if not fuente:
            err.append('regla sin fuente: ' + k)
    for t, (fmt, src, desc) in TOKENS_DOCX.items():
        if fmt not in FMT or not re.fullmatch(r'[a-z0-9_]+', t):
            err.append('token mal formado: ' + t)
        if src[0] == 'fila' and src[3] not in 'BCDEFGH':
            err.append('token con columna rara: ' + t)
    for plan in PLANES:
        fij = textos_fijos_docx(plan)
        for i, v in fij.items():
            if restos_espanol(v) or RX_MARCA_ES.search(v):
                err.append('DOCX_FIJOS %s[%d] con español o marca ES: %r' % (plan, i, v[:50]))
    # formatear
    pruebas = [(254800, 'usd', '$254,800'), (14, 'usd2', '$14.00'), (0.299, 'pct1', '29.9%'), (0.08, 'pct0', '8%'),
               (1.8412, 'x2', '1.84×'), (41.02, 'ceil', '42'), (41.0, 'ceil', '41'), (2.84, 'anos1', '2.8 years'),
               ('Over 3 years', 'anos1', 'more than 3 years'), (-1460.4, 'usd', '-$1,460'), (4.25, 'dec2', '4.25')]
    for v, f, esp in pruebas:
        if formatear(v, f) != esp:
            err.append('formatear(%r, %s) = %r ≠ %r' % (v, f, formatear(v, f), esp))
    return err


def _celdas(sqref):
    from openpyxl.utils import range_boundaries, get_column_letter
    s = set()
    for rng in str(sqref).split():
        c1, r1, c2, r2 = range_boundaries(rng)
        for r in range(r1, r2 + 1):
            for c in range(c1, c2 + 1):
                s.add(get_column_letter(c) + str(r))
    return s


def celdas_formulas_en():
    """(corto, hoja ES, celda) cubiertas por FORMULAS_EN (expande rangos de una fila)."""
    out = {}
    for (c, h, rg), v in FORMULAS_EN.items():
        for x in _celdas(rg):
            out[(c, h, x)] = v
    return out


def cruzar_censo(p):
    censo = json.load(open(p, encoding='utf-8'))
    err = []
    cubiertas = celdas_formulas_en()
    for f, info in censo['libros'].items():
        c = info['corto']
        if f not in FICHEROS:
            err.append('libro sin nombre EN: ' + f)
        if HOJAS_ES_POR_LIBRO.get(c) != info['hojas']:
            err.append('HOJAS_ES_POR_LIBRO no coincide con el censo en %s: %s' % (c, info['hojas']))
        for h, hd in info['hojas_detalle'].items():
            for dv in hd['dv']:
                f1 = dv['formula1'] or ''
                if f1.startswith('"'):
                    for it in f1.strip('"').split(','):
                        if necesita_clave(it) and it not in CLAVES:
                            err.append('ítem de DV sin EN: %s / %s / %r' % (c, h, it))
            # literales de fórmula: CLAVES o la celda entera en FORMULAS_EN
            for lit, celdas in hd['literales'].items():
                if not necesita_clave(lit) or lit in CLAVES:
                    continue
                sin = [x for x in celdas if (c, h, x) not in cubiertas]
                if sin:
                    err.append('literal sin EN: %s / %s / %r en %s' % (c, h, lit[:50], ','.join(sin[:4])))
            for lit in hd['literales_cf']:
                if necesita_clave(lit) and lit not in CLAVES:
                    err.append('literal de CF sin EN: %s / %s / %r' % (c, h, lit))
            # FORMULAS_EN: la celda tiene fórmula en el ES
            for (c2, h2, x), v in cubiertas.items():
                if c2 == c and h2 == h and x not in hd['formulas_texto']:
                    err.append('FORMULAS_EN cita una celda sin fórmula de texto en el ES: %s / %s / %s' % (c, h, x))
            # CF: tokens EN inyectivos y ningún literal cambia de color
            por_rango = OrderedDict()
            for cf in hd['cf']:
                por_rango.setdefault(cf['sqref'], []).extend(cf['tokens'])
            for rng, toks in por_rango.items():
                toks_en = [CLAVES.get(t, t) for t in toks]
                if len(set(t.lower() for t in toks_en)) != len(set(t.lower() for t in toks)):
                    err.append('tokens de CF EN no inyectivos en %s / %s / %s' % (c, h, rng))
                for t_es, t_en in zip(toks, toks_en):
                    for t2_es, t2_en in zip(toks, toks_en):
                        if t_es != t2_es and t_en.lower() in t2_en.lower() and t_es.lower() not in t2_es.lower():
                            err.append('token EN %r contenido en %r (%s / %s / %s)' % (t_en, t2_en, c, h, rng))
                cs = _celdas(rng)
                for lit, celdas in hd['literales'].items():
                    if not (cs & set(celdas)) or not necesita_clave(lit) or lit not in CLAVES:
                        continue
                    esperado = {CLAVES.get(t, t).lower() for t in toks if t.lower() in lit.lower()}
                    real = {t.lower() for t in toks_en if t.lower() in CLAVES[lit].lower()}
                    if esperado != real:
                        err.append('el literal %r cambia de color en %s / %s / %s' % (CLAVES[lit], c, h, rng))
    # PARCHES y VALORES_EN contra el censo
    for (c, h, x), (viejo, nuevo) in PARCHES_FORMULA.items():
        info = next(i for i in censo['libros'].values() if i['corto'] == c)
        fx = info['hojas_detalle'][h]['formulas'].get(x)
        if not fx or viejo not in fx:
            err.append('PARCHE E2 que no casa con el ES: %s / %s / %s' % (c, h, x))
    for (c, h, x), (v, rot) in VALORES_EN.items():
        info = next(i for i in censo['libros'].values() if i['corto'] == c)
        hd = info['hojas_detalle'][h]
        if x not in hd['verdes_celdas']:
            err.append('VALORES_EN sobre una celda que no es verde en el ES: %s / %s / %s' % (c, h, x))
        if rot:
            fila = re.sub(r'[A-Z]+', '', x)
            et = hd['etiquetas'].get(fila) or ''
            if not et.startswith(rot):
                err.append('VALORES_EN %s / %s / %s: la fila dice %r y se esperaba %r' % (c, h, x, et[:40], rot))
    # tokens resueltos
    for plan in PLANES:
        res = censo.get('tokens_docx', {}).get(plan, {})
        for t, (fmt, src, desc) in tokens_de(plan).items():
            if src[0] == 'fila' and t not in res:
                err.append('token sin resolver en %s: %s' % (plan, t))
    return err


# ==========================================================================
# Conjunto restaurante + panadería: sustituye los datos de este módulo (ver CONJUNTO arriba)
# ==========================================================================
if CONJUNTO == 'restpan':
    import sys as _sys
    _sys.path.insert(0, os.path.join(os.path.dirname(AQUI), 'business-plans-en-2'))
    import mapas_rp as _rp                                                         # noqa: E402
    _rp.instalar(_sys.modules[__name__])
elif CONJUNTO != 'ftcaf':
    raise SystemExit('BP_CONJUNTO desconocido: %r (ftcaf | restpan)' % CONJUNTO)


if __name__ == '__main__':
    e = autotest()
    pc = os.path.join(DATOS, 'censo_es.json')
    if os.path.exists(pc):
        e += cruzar_censo(pc)
    for x in e:
        print('ERROR', x)
    print('mapas.py: {} libros · {} pestañas · {} claves · {} fijos · {} fórmulas EN · {} valores · {} tokens · '
          '{} errores'.format(len(LIBROS), sum(len(d) for d in HOJAS_POR_LIBRO.values()), len(CLAVES), len(FIJOS),
                              len(FORMULAS_EN), len(VALORES_EN), len(TOKENS_DOCX), len(e)))
    raise SystemExit(1 if e else 0)
