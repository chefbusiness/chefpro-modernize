#!/usr/bin/env python3
"""
extraer_textos.py — Food Truck + Coffee Shop Business Plan Kit (EN) · cimientos de la F2-EN (xlsx).

Calcado de `financial-kit/extraer_textos.py` (sesión Claude Code, 4-oct-2026). Lee los 4 xlsx PUBLICADOS de
`astro-site/public/dl/plan-negocio-food-truck/` y `…/plan-negocio-cafeteria/` (SOLO LECTURA) y escribe, junto a
este script (o en --salida):

  · censo_es.json   por libro y hoja, lo que contarán los gates de la F2-EN: fórmulas (todas, con huella sha1 y
                    PATRONES), referencias a hojas, literales en fórmulas, CF y DV, DV, CF (con tokens y relleno),
                    merges, protección, paneles, papel, cabeceras/pies, formatos (con €), celdas verdes y
                    desbloqueadas (con su lista), rótulos de la col. A, TAREAS por fase (checklists), docProps y la
                    resolución de los TOKENS del docx (mapas.TOKENS_DOCX → hoja/celda en cada plan, con su valor ES
                    en caché como comprobación de tipo).
  · textos_es.json  cadenas de texto únicas con su contexto (`donde`), TRATAMIENTO, PISTAS de la SPEC, LÍMITES y
                    GRUPO de traducción: GM (comunes FT + CAF), GFT (solo food truck), GCAF (solo cafetería) para
                    los subagentes; GX = NO se traduce (pestañas, CLAVES, FIJOS, POR_CELDA, versión, pie, docProps).

    python3 extraer_textos.py [--origen-repo DIR] [--salida DIR] [--md]

No traduce nada (regla 1bis). Térmica: un proceso, en serie, 4 libros, < 10 s. Ejecutar en el VPS (D27).

Tratamientos (la cadena toma el más «fuerte»: traducir > localizar > mapa-hoja > mapa-clave > regenerar):
  traducir    texto libre → subagente, con pistas.
  localizar   celda de una ZONA de mapas.ZONAS_CELDA (tablas D18/D19 de Instrucciones): UN id por celda, también
              las celdas sin letras («28-33 %»), porque el EN depende de la fila.
  mapa-hoja   texto igual a una pestaña del libro → mapas.HOJAS_POR_LIBRO.
  mapa-clave  ítem de DV o literal de fórmula/CF del libro (o celda igual a uno) → mapas.CLAVES.
  regenerar   lo escribe aplicar_en.py: versión, pie, docProps, mapas.FIJOS, mapas.POR_CELDA.
Los literales de fórmula NO son cadenas de celda: viven en censo_es.json y los cubre mapas.FORMULAS_EN / CLAVES.
"""
import argparse
import datetime
import hashlib
import json
import os
import re
import subprocess
import sys
from collections import OrderedDict, defaultdict

import openpyxl
from openpyxl.utils import get_column_letter, range_boundaries

AQUI = os.path.dirname(os.path.abspath(__file__))
REPO = os.path.abspath(os.path.join(AQUI, '..', '..', '..'))
sys.path.insert(0, AQUI)
import mapas                                                     # noqa: E402

VERDE = 'E8F5E9'
INVARIABLES = {'OK', 'N/A', 'AI Chef Pro', 'EBITDA', 'Marketing', 'DSCR', 'Food Truck', 'Break-even'}
RX_VERSION = re.compile(r'^Versi[óo]n \d+\.\d+ · ')
RX_LITERAL = re.compile(r'"((?:[^"]|"")*)"')
RX_REF_HOJA = re.compile(r"(?:'((?:[^']|'')+)'|([^\W\d][\w.]*))!(?=\$?[A-Za-z]{1,3}\$?\d|\$?[A-Za-z]{1,3}:)")
RX_CITA = re.compile(r"«([^»]{1,40})»|'([^']{1,40})'|\"([^\"]{1,40})\"")
RX_NORMA = re.compile(r'\bIVA\b|IRPF|Seguridad Social|\bSS\b|RETA|[Mm]od(?:elo|\.) ?\d{3}|AEAT|Hacienda|\bICO\b|'
                      r'ENISA|\bLIS\b|\bLIVA\b|\bIAE\b|\bSMI\b|\bCCAA\b|[Ii]mpuesto (?:de|sobre) Sociedades|\bSL\b|'
                      r'[Nn]otar|Registro Mercantil|\bNIF\b|\bCIF\b|OEPM|RGSEAA|SGAE|AGEDI|Veri\*factu|[Cc]onvenio|'
                      r'[Gg]estor|\bTIN\b|\bTAE\b|[Aa]utónom|autonómic|\bart\.|[Cc]arencia|[Aa]val\b|[Ff]ianza|'
                      r'\bITV\b|ROESB|APPCC|\bDDD\b|\bPRL\b|RGPD|LOPDGDD|AEPD|FNMT|\bBOE\b|\bRD\b|\bLey\b|Estatuto|'
                      r'Ayuntamiento|[Mm]unicip|Comunidad Autónoma|registro sanitario|Sanidad')
RX_MONEDA = re.compile(r'€|\bEUR\b|euros?\b')
RX_M2 = re.compile(r'm²')
RX_COMPARA = re.compile(r'[<>]=?\s*-?\d')

PISTAS = [
    (RX_NORMA.pattern, 'D9-D13 y §4.2: trámite/ley/organismo español → su equivalente US (SPEC §4.2, research §5); '
                       'nota UK breve SOLO donde la SPEC la pone. Nunca se traduce un modelo, ley u organismo español: '
                       'se sustituye o se quita'),
    (RX_MONEDA.pattern, 'D8: sin símbolo de moneda; «(€)» fuera; «€/mes» → «per month»; importes en texto sin símbolo'),
    (r'm²|metros cuadrados', 'D14: m² → sq ft (1 m² = 10.764 sq ft)'),
    (r'litros|\bl\b|kg\b', 'D14: litros → galones (exacto + «check your county»); 3.500 kg → regla CDL (GVWR ≥ 26,001 lb)'),
    (r'[Ss]in IVA|SIN IVA|[Cc]on IVA|IVA incl|IVA incluido', 'D9: «sin IVA» → «excl. sales tax»; «con IVA» → «incl. sales '
                                                            'tax»; P&L sin impuesto, tesorería con lo cobrado'),
    (r'IVA soportado|recuperable', 'D9: US: el sales tax pagado es coste (input tax recuperable = 0); UK: 20 % si está '
                                   'registrado en VAT'),
    (r'303|trimestr', 'D9: liquidación trimestral (abril, julio, octubre, enero) = «quarterly filer»; mensual → nota'),
    (r'[Ii]mpuesto (?:de|sobre) Sociedades|[Bb]ases negativas|nueva creación',
     'D10: «effective income tax rate» 25 % (21 % federal C corp + estatal; pass-through 0 %); bases negativas → '
     '«net operating losses carried forward» (federal: hasta el 80 % de la base; el modelo compensa el 100 %)'),
    (r'Seguridad Social|\bSS\b|pagas|paga extra|SMI|[Cc]onvenio|Estatuto|vacaciones',
     'D11: «employer payroll taxes» 10 % (FICA + FUTA/SUTA + workers\' comp); 12 pay periods (las extras solo si > 12); '
     'suelo = el mayor entre el mínimo federal ($7.25/h × 2,080 h) y el estatal/local; 40 h (FLSA); 50 semanas; '
     'propinas no modeladas'),
    (r'[Aa]mortizaci', 'D12: amortización del inmovilizado → «depreciation» (lineal de libros; MACRS / Section 179 solo '
                       'en nota); amortización del préstamo → «principal repayment»; cuadro → «amortization schedule»'),
    (r'ICO|ENISA|[Ss]ubvenci|[Bb]usiness angels|[Pp]réstamo|banco|[Ee]ntidad',
     'D13: banco → «lender»; SBA 7(a) a través del banco; SBA Microloan (hasta $50,000); tipo 10 % «not APR»; '
     'aportación propia: «lenders usually expect an owner equity injection (often around 10% for SBA start-ups; many '
     'want 20-30%)»; formato «lender-ready», nunca «approval guaranteed»'),
    (r'[Cc]arencia', 'D13/D17: carencia → «interest-only period» (0 por defecto en la EN)'),
    (r'DSCR', 'D13: DSCR límite 1.15× y objetivo 1.25×; NUNCA «SBA minimum»'),
    (r'[Aa]parcamiento|base del vehículo|nave', 'Glosario: aparcamiento y base del vehículo → «commissary & truck parking»'),
    (r'[Vv]enta ambulante|vía pública', 'Glosario: venta ambulante → «mobile vending»; vía pública → «public right-of-way '
                                        '/ vending zones»'),
    (r'[Gg]estor|gestoría', 'Glosario: gestor/gestoría → «accountant / bookkeeper»; gestor de residuos → «licensed waste '
                            '/ grease hauler»'),
    (r'[Cc]lientes?|[Tt]icket medio|PVP', 'Glosario: clientes/día → «customers per day»; ticket medio → «average check»; '
                                          'PVP con IVA → «menu price incl. sales tax»'),
    (r'TPV|[Dd]atáfono|bizum', 'Glosario: TPV → «POS»; datáfono → «card terminal»; bizum → «mobile wallet / contactless»'),
    (r'v1\.1|versión anterior|[Ff]ichero v|la versión 1|plan anterior|antes', 'D18: la EN nace en 2.2: explica el porqué '
                                                                              'sin historia de versiones'),
    (r'España|español|Madrid|Barcelona|Valencia|Bilbao|San Miguel|Boqueria|costa',
     'Mercado US (ejemplo): sin ciudades ni mercados españoles; «your area», «your city»'),
    (r'Revisar →|[Dd]esproteger|Archivo →', 'Rutas de Excel en inglés: «Review → Unprotect Sheet»'),
    (r"[«»]|'[A-ZÁÉÍÓÚ0-9]", 'D6: pestañas y valores citados entre comillas dobles rectas con su nombre EN '
                             '("0. Assumptions", "Staffing")'),
    (r'\d,\d', 'Formato US: punto decimal y coma de miles («0.35 = 35%»)'),
    (r'«Sí»|«No»|"Sí"|Sí/No', 'D7: los valores de la lista son «Yes» / «No»'),
    (r'CUMPLE|REVISAR', 'D7: «OK» / «REVIEW»'),
    (r'verificad|vigente', 'A3: el EN no afirma «verified» en 50 estados: «a starting point; requirements vary by state, '
                           'county and city»'),
    (r'alérgen|manipulador', 'Research §5: Certified Food Protection Manager + food handler cards; 9 alérgenos US '
                             '(sésamo desde 2023; UK 14)'),
    (r'música|SGAE|AGEDI', 'Research §5: ASCAP / BMI / SESAC si se pone música'),
    (r'Glovo|Uber|TheFork|Google My Business|TripAdvisor|RRSS',
     'SPEC §4.2: DoorDash / Uber Eats, Resy / OpenTable, Google Business Profile, Yelp; apps de food trucks (Roaming '
     'Hunger, StreetFoodFinder); RRSS → social media'),
]


# --------------------------------------------------------------------------
# utilidades
# --------------------------------------------------------------------------
def sha256(path):
    h = hashlib.sha256()
    with open(path, 'rb') as fh:
        for b in iter(lambda: fh.read(65536), b''):
            h.update(b)
    return h.hexdigest()


def tiene_letra(s):
    return any(ch.isalpha() for ch in s)


def literales(formula):
    return [m.group(1).replace('""', '"') for m in RX_LITERAL.finditer(formula or '')]


def refs_hoja(formula):
    sin_lit = RX_LITERAL.sub('""', formula or '')
    return [(m.group(1) or m.group(2)).replace("''", "'") for m in RX_REF_HOJA.finditer(sin_lit)]


def celdas_de(sqref):
    s = set()
    for rng in str(sqref).split():
        c1, r1, c2, r2 = range_boundaries(rng)
        for r in range(r1, r2 + 1):
            for c in range(c1, c2 + 1):
                s.add((r, c))
    return s


def items_dv(f1):
    return f1.strip('"').split(',') if f1 and f1.startswith('"') else []


def claves_de_libro(wb):
    s = set()
    for ws in wb.worksheets:
        for dv in ws.data_validations.dataValidation:
            s.update(items_dv(dv.formula1))
            if dv.type == 'custom':
                s.update(literales(dv.formula1))
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.data_type == 'f':
                    s.update(literales(c.value))
        for rng, reglas in ws.conditional_formatting._cf_rules.items():
            for rg in reglas:
                for f in rg.formula or []:
                    s.update(literales(f))
    return s


def zonas_de(corto, ws):
    """Zonas celda a celda resueltas: lista de (f1, f2, cols, motivo). Las de la CAF se buscan por cabecera."""
    out = []
    for f1, f2, cols, motivo in mapas.ZONAS_CELDA.get((corto, ws.title), []):
        if f1 is None:
            cab = mapas.CABECERA_ZONA[motivo]
            fila_cab = next((c.row for c in ws['A'] if c.value == cab), None)
            if fila_cab is None:
                raise SystemExit('zona %s: no encuentro la cabecera %r en %s/%s' % (motivo, cab, corto, ws.title))
            f1 = fila_cab + 2                          # cabecera de la tabla en fila_cab + 1
            f2 = f1
            while ws.cell(f2 + 1, 1).value not in (None, ''):
                f2 += 1
        out.append((f1, f2, cols, motivo))
    return out


def en_zona(zonas, cel):
    for f1, f2, cols, motivo in zonas:
        if f1 <= cel.row <= f2 and get_column_letter(cel.column) in cols:
            return motivo
    return None


def clasificar_celda(corto, ws, cel, texto, claves_libro, zonas):
    if (corto, ws.title, cel.coordinate) in mapas.POR_CELDA:
        return 'regenerar', 'rótulo fijado (mapas.POR_CELDA)'
    if RX_VERSION.match(texto):
        return 'regenerar', 'D20 línea de versión'
    if texto in mapas.FIJOS:
        return 'regenerar', 'decisión de la SPEC (mapas.FIJOS)'
    z = en_zona(zonas, cel)
    if z:
        return 'localizar', z
    if texto in mapas.HOJAS_POR_LIBRO[corto]:
        return 'mapa-hoja', 'celda igual a una pestaña'
    if texto in claves_libro and texto in mapas.CLAVES:
        return 'mapa-clave', 'valor de DV/literal del libro'
    return 'traducir', None


# --------------------------------------------------------------------------
# censo de un libro
# --------------------------------------------------------------------------
def censar_libro(path, fichero):
    corto = mapas.corto_de(fichero)
    wb = openpyxl.load_workbook(path)
    wbv = openpyxl.load_workbook(path, data_only=True)
    hojas = list(wb.sheetnames)
    ocurr = []
    info = OrderedDict()
    info['corto'] = corto
    info['plan'] = mapas.PLAN_DE[corto]
    info['tipo'] = mapas.TIPO_DE[corto]
    info['sha256'] = sha256(path)
    info['nombre_en'] = mapas.FICHEROS[fichero]
    info['titulo_en'] = mapas.TITULOS[mapas.FICHEROS[fichero]]
    info['hojas'] = hojas
    info['hojas_en'] = [mapas.HOJAS_POR_LIBRO[corto].get(h) for h in hojas]
    info['nombres_definidos'] = sorted(getattr(wb.defined_names, 'keys', lambda: [])())
    pr = wb.properties
    info['docprops'] = OrderedDict((k, getattr(pr, k)) for k in
                                   ('creator', 'title', 'subject', 'keywords', 'description', 'category',
                                    'lastModifiedBy'))
    info['hojas_detalle'] = OrderedDict()
    claves_libro = claves_de_libro(wb)
    for ws in wb.worksheets:
        ocurr.append((ws.title, {'f': corto, 'h': ws.title, 't': 'hoja', 'tr': 'mapa-hoja', 'det': 'nombre de pestaña'}))
        wsv = wbv[ws.title]
        zonas = zonas_de(corto, ws)
        hd = OrderedDict()
        hd['nombre_en'] = mapas.HOJAS_POR_LIBRO[corto].get(ws.title)
        hd['dimension'] = ws.dimensions
        hd['max_fila'] = ws.max_row
        hd['max_col'] = ws.max_column
        hd['zonas_celda'] = [OrderedDict([('filas', '%d-%d' % (a, b)), ('cols', c), ('motivo', m)])
                             for a, b, c, m in zonas]
        n_txt = n_num = n_fx = fx_con_hoja = 0
        refs = defaultdict(int)
        lits = defaultdict(list)
        formatos = defaultdict(int)
        numeros, m2, norma, moneda, fmt_euro = [], [], [], [], []
        verdes, desbl = [], 0
        verdes_desbl = 0
        patrones = OrderedDict()
        formulas = OrderedDict()
        formulas_texto = OrderedDict()
        etiquetas = OrderedDict()
        huella = hashlib.sha1()
        for row in ws.iter_rows():
            for cel in row:
                if cel.__class__.__name__ == 'MergedCell':
                    continue
                v = cel.value
                fmt = cel.number_format
                if fmt and fmt != 'General':
                    formatos[fmt] += 1
                if fmt and '€' in fmt:
                    fmt_euro.append(cel.coordinate)
                fill = cel.fill.fgColor.rgb if cel.fill is not None and cel.fill.fgColor is not None else None
                bloq = cel.protection is None or cel.protection.locked is not False
                if not bloq:
                    desbl += 1
                if isinstance(fill, str) and fill.upper().endswith(VERDE):
                    verdes.append(cel.coordinate)
                    if not bloq:
                        verdes_desbl += 1
                if v is None:
                    continue
                if cel.column == 1 and isinstance(v, str) and cel.data_type != 'f':
                    etiquetas[str(cel.row)] = v
                if isinstance(v, str) and cel.data_type == 'f':
                    n_fx += 1
                    formulas[cel.coordinate] = v
                    huella.update('{}={}\n'.format(cel.coordinate, v).encode('utf-8'))
                    lt = literales(v)
                    for lit in lt:
                        lits[lit].append(cel.coordinate)
                    if any(mapas.necesita_clave(x) and x not in mapas.CLAVES for x in lt):
                        formulas_texto[cel.coordinate] = v
                    rh = refs_hoja(v)
                    if rh:
                        fx_con_hoja += 1
                    for nom in rh:
                        refs[nom] += 1
                    p = patron(v)
                    if p not in patrones:
                        patrones[p] = OrderedDict([('patron', p), ('n', 0), ('celdas', []),
                                                   ('cifras', bool(RX_COMPARA.search(RX_LITERAL.sub('""', v))))])
                    patrones[p]['n'] += 1
                    patrones[p]['celdas'].append(cel.coordinate)
                    continue
                if isinstance(v, str):
                    n_txt += 1
                    if RX_M2.search(v):
                        m2.append(cel.coordinate)
                    if RX_NORMA.search(v):
                        norma.append(cel.coordinate)
                    if RX_MONEDA.search(v):
                        moneda.append(cel.coordinate)
                    tr, det = clasificar_celda(corto, ws, cel, v, claves_libro, zonas)
                    d = {'f': corto, 'h': ws.title, 'c': cel.coordinate, 't': 'celda', 'tr': tr}
                    if det:
                        d['det'] = det
                    ocurr.append((v, d))
                else:
                    n_num += 1
                    if ws.title != 'Instrucciones':
                        val = v.isoformat() if hasattr(v, 'isoformat') else v
                        numeros.append([cel.coordinate, val, fmt])
                    elif en_zona(zonas, cel):         # número escrito en una zona celda a celda (raro): se localiza
                        ocurr.append((str(v), {'f': corto, 'h': ws.title, 'c': cel.coordinate, 't': 'celda',
                                               'tr': 'localizar', 'det': en_zona(zonas, cel)}))
        for p in patrones.values():
            cs = p['celdas']
            p['celdas'] = cs[0] + ('..' + cs[-1] if len(cs) > 1 else '')
        hd['celdas_texto'] = n_txt
        hd['celdas_numero'] = n_num
        hd['celdas_formula'] = n_fx
        hd['formulas_con_hoja'] = fx_con_hoja
        hd['refs_hoja'] = OrderedDict(sorted(refs.items()))
        hd['literales'] = OrderedDict((k, lits[k]) for k in sorted(lits))
        hd['literales_recuento'] = OrderedDict((k, len(lits[k])) for k in sorted(lits))
        hd['huella_formulas_sha1'] = huella.hexdigest() if n_fx else None
        hd['patrones_formula'] = list(patrones.values())
        hd['formulas'] = formulas
        hd['formulas_texto'] = formulas_texto
        hd['etiquetas'] = etiquetas
        hd['numeros'] = numeros
        hd['formatos'] = OrderedDict(sorted(formatos.items(), key=lambda kv: -kv[1]))
        hd['formatos_euro'] = fmt_euro
        hd['textos_m2'] = m2
        hd['textos_normativa_es'] = norma
        hd['textos_moneda'] = moneda
        hd['verdes'] = len(verdes)
        hd['verdes_celdas'] = verdes
        hd['verdes_desbloqueadas'] = verdes_desbl
        hd['desbloqueadas'] = desbl
        hd['columnas_ocultas'] = sorted(k for k, d in ws.column_dimensions.items() if d.hidden)
        hd['filas_ocultas'] = sorted(k for k, d in ws.row_dimensions.items() if d.hidden)
        hd['merges'] = sorted(str(m) for m in ws.merged_cells.ranges)
        hd['paneles'] = ws.freeze_panes
        p = ws.protection
        hd['proteccion'] = OrderedDict([('hoja', bool(p.sheet)), ('password', bool(p.password)),
                                        ('formatColumns', p.formatColumns), ('formatRows', p.formatRows),
                                        ('insertRows', p.insertRows), ('deleteRows', p.deleteRows),
                                        ('sort', p.sort), ('autoFilter', p.autoFilter)])
        hd['autofiltro'] = ws.auto_filter.ref
        hd['area_impresion'] = ws.print_area
        hd['titulos_impresion'] = OrderedDict([('filas', ws.print_title_rows), ('columnas', ws.print_title_cols)])
        ps = ws.page_setup
        hd['papel'] = OrderedDict([('paperSize', ps.paperSize), ('orientation', ps.orientation),
                                   ('fitToWidth', ps.fitToWidth), ('fitToHeight', ps.fitToHeight),
                                   ('fitToPage', bool(ws.sheet_properties.pageSetUpPr and
                                                      ws.sheet_properties.pageSetUpPr.fitToPage))])
        hf = OrderedDict()
        for nombre in ('oddHeader', 'oddFooter', 'evenHeader', 'evenFooter', 'firstHeader', 'firstFooter'):
            obj = getattr(ws, nombre)
            for parte in ('left', 'center', 'right'):
                txt = getattr(obj, parte).text
                if txt:
                    hf['{}.{}'.format(nombre, parte)] = txt
                    ocurr.append((txt, {'f': corto, 'h': ws.title, 't': 'cabecera-pie',
                                        'c': '{}.{}'.format(nombre, parte), 'tr': 'regenerar',
                                        'det': 'D8 pie «%s»' % mapas.FOOTER_EN}))
        hd['cabeceras_pies'] = hf
        dvs = []
        for dv in ws.data_validations.dataValidation:
            dvs.append(OrderedDict([
                ('sqref', str(dv.sqref)), ('type', dv.type), ('operator', dv.operator),
                ('formula1', dv.formula1), ('formula2', dv.formula2), ('allow_blank', dv.allow_blank),
                ('showErrorMessage', dv.showErrorMessage), ('showInputMessage', dv.showInputMessage),
                ('errorStyle', dv.errorStyle), ('errorTitle', dv.errorTitle), ('error', dv.error),
                ('promptTitle', dv.promptTitle), ('prompt', dv.prompt),
                ('n_celdas', len(celdas_de(dv.sqref))),
                ('cita_hojas', sorted(set(refs_hoja(dv.formula1 or '')))),
                ('literales', literales(dv.formula1) if dv.type == 'custom' else []),
            ]))
            base = {'f': corto, 'h': ws.title, 'c': str(dv.sqref)}
            for campo, t in (('errorTitle', 'dv-error-titulo'), ('error', 'dv-error'),
                             ('promptTitle', 'dv-prompt-titulo'), ('prompt', 'dv-prompt')):
                txt = getattr(dv, campo)
                if txt:
                    ocurr.append((txt, dict(base, t=t, tr='traducir')))
            for it in items_dv(dv.formula1):
                if it:
                    ocurr.append((it, dict(base, t='dv-item', tr='mapa-clave', det='mapas.CLAVES')))
        hd['dv'] = dvs
        cfs = []
        lits_cf = defaultdict(int)
        for rng, reglas in ws.conditional_formatting._cf_rules.items():
            for rg in reglas:
                fs = list(rg.formula or [])
                cfs.append(OrderedDict([
                    ('sqref', str(rng.sqref)), ('type', rg.type), ('operator', rg.operator),
                    ('text', getattr(rg, 'text', None)), ('formula', fs), ('priority', rg.priority),
                    ('stopIfTrue', rg.stopIfTrue), ('cita_hojas', sorted({h for f in fs for h in refs_hoja(f)})),
                    ('tokens', [t for f in fs for t in literales(f)]),
                    ('fill', getattr(getattr(getattr(rg.dxf, 'fill', None), 'fgColor', None), 'rgb', None)
                     if rg.dxf is not None else None),
                ]))
                for f in fs:
                    for lit in literales(f):
                        lits_cf[lit] += 1
        hd['cf'] = cfs
        hd['literales_cf'] = OrderedDict(sorted(lits_cf.items()))
        # checklists: tareas = celdas de la DV de lista de la columna de OK (✓/☐/N/A o ✓/—/N/A); en el molde de
        # filas-cabecera (RESTC) las cabeceras de fase no son tareas
        if info['tipo'] == 'check' and ws.title != 'Instrucciones':
            col_ok = mapas.COL_OK.get(corto, 1)
            cabs = set(mapas.FASES_CABECERA[corto][1]) if corto in mapas.FASES_CABECERA else set()
            tareas = sorted({r for dv in ws.data_validations.dataValidation if (dv.formula1 or '').startswith('"✓')
                             for (r, c) in celdas_de(dv.sqref) if c == col_ok and r not in cabs})
            hd['tareas'] = len(tareas)
            hd['tareas_filas'] = '%d-%d' % (tareas[0], tareas[-1]) if tareas else None
            hd['tareas_con_texto'] = sum(1 for r in tareas if isinstance(ws.cell(r, 2).value, str))
        info['hojas_detalle'][ws.title] = hd
        info.setdefault('_wsv', {})[ws.title] = wsv
    if info['tipo'] == 'check':
        info['tareas_por_fase'] = mapas.contar_tareas(corto, wb, en=False)
    for campo, v in info['docprops'].items():
        if not isinstance(v, str) or not v or campo == 'creator':
            continue
        ocurr.append((v, {'f': corto, 't': 'docprops', 'c': campo, 'tr': 'regenerar', 'det': 'D20 docProps'}))

    hdv = list(info['hojas_detalle'].values())
    lit_tot, litcf_tot = defaultdict(int), defaultdict(int)
    for x in hdv:
        for k, v in x['literales_recuento'].items():
            lit_tot[k] += v
        for k, v in x['literales_cf'].items():
            litcf_tot[k] += v
    info['totales'] = OrderedDict([
        ('hojas', len(hojas)),
        ('celdas_texto', sum(x['celdas_texto'] for x in hdv)),
        ('celdas_formula', sum(x['celdas_formula'] for x in hdv)),
        ('formulas_con_hoja', sum(x['formulas_con_hoja'] for x in hdv)),
        ('formulas_texto', sum(len(x['formulas_texto']) for x in hdv)),
        ('patrones_formula', sum(len(x['patrones_formula']) for x in hdv)),
        ('dv', sum(len(x['dv']) for x in hdv)),
        ('dv_lista_literal', sum(1 for x in hdv for d in x['dv'] if (d['formula1'] or '').startswith('"'))),
        ('dv_celdas', sum(d['n_celdas'] for x in hdv for d in x['dv'])),
        ('cf', sum(len(x['cf']) for x in hdv)),
        ('merges', sum(len(x['merges']) for x in hdv)),
        ('hojas_protegidas', sum(1 for x in hdv if x['proteccion']['hoja'])),
        ('hojas_con_password', sum(1 for x in hdv if x['proteccion']['password'])),
        ('paneles', sum(1 for x in hdv if x['paneles'])),
        ('numeros_ejemplo', sum(len(x['numeros']) for x in hdv)),
        ('formatos_euro', sum(len(x['formatos_euro']) for x in hdv)),
        ('textos_m2', sum(len(x['textos_m2']) for x in hdv)),
        ('textos_normativa_es', sum(len(x['textos_normativa_es']) for x in hdv)),
        ('textos_moneda', sum(len(x['textos_moneda']) for x in hdv)),
        ('verdes', sum(x['verdes'] for x in hdv)),
        ('verdes_desbloqueadas', sum(x['verdes_desbloqueadas'] for x in hdv)),
        ('desbloqueadas', sum(x['desbloqueadas'] for x in hdv)),
        ('tareas', sum(x.get('tareas', 0) for x in hdv)),
        ('papel', sorted({x['papel']['paperSize'] for x in hdv if x['papel']['paperSize'] is not None})),
        ('literales', OrderedDict(sorted(lit_tot.items()))),
        ('literales_distintos', len(lit_tot)),
        ('literales_texto_distintos', sum(1 for k in lit_tot if mapas.necesita_clave(k))),
        ('literales_cf', OrderedDict(sorted(litcf_tot.items()))),
    ])
    return info, ocurr


def patron(formula):
    """Normaliza las filas de las referencias relativas a {r} (las absolutas $X$n se respetan)."""
    return re.sub(r'(?<![A-Za-z_$])(\$?[A-Z]{1,3})(?<!\$)(\d+)(?![\d(])', lambda m: m.group(1) + '{r}', formula)


def resolver_tokens(censo):
    """mapas.TOKENS_DOCX ('fila', hoja, rótulo, col) → celda en cada plan + valor ES en caché (comprobación)."""
    out = OrderedDict()
    for plan in mapas.PLANES:
        corto = mapas.corto_plan(plan)
        f = mapas.LIBROS[corto][1]
        info = censo['libros'][f]
        res = OrderedDict()
        for t, (fmt, src, desc) in mapas.tokens_de(plan).items():
            if src[0] != 'fila':
                continue
            _, hoja, rot, col = src
            rot = rot[plan] if isinstance(rot, dict) else rot
            hoja = mapas.hoja_es(corto, hoja)                   # nombre canónico → pestaña real del libro
            filas = [r for r, et in info['hojas_detalle'][hoja]['etiquetas'].items() if et == rot]
            if len(filas) != 1:
                raise SystemExit('token %s/%s: el rótulo %r aparece %d veces en %s' % (plan, t, rot, len(filas), hoja))
            celda = '%s%s' % (col, filas[0])
            v = info['_wsv'][hoja][celda].value
            res[t] = OrderedDict([('hoja_es', hoja), ('hoja_en', mapas.hoja_en(corto, hoja)), ('celda', celda),
                                  ('fmt', fmt), ('valor_es', v)])
            if not isinstance(v, (int, float)) and not (fmt == 'anos1' and isinstance(v, str)):
                raise SystemExit('token %s/%s: %s!%s no es un número en el ES (%r)' % (plan, t, hoja, celda, v))
        out[plan] = res
    return out


# --------------------------------------------------------------------------
# cadenas, grupos y pistas
# --------------------------------------------------------------------------
PRIORIDAD = {'traducir': 0, 'localizar': 1, 'mapa-hoja': 2, 'mapa-clave': 3, 'regenerar': 4}
LIMITE = {'dv-error-titulo': 32, 'dv-prompt-titulo': 32, 'dv-error': 255, 'dv-prompt': 255}


def pistas(texto):
    return [nota for rx, nota in PISTAS + list(mapas.PISTAS_EXTRA) if re.search(rx, texto)]


def citas_hoja(texto, todas_hojas, cortos):
    out = []
    for m in RX_CITA.finditer(texto):
        span = next(g for g in m.groups() if g is not None)
        if span in todas_hojas and span not in [c['es'] for c in out]:
            en = sorted({mapas.HOJAS_POR_LIBRO[c].get(span) for c in cortos if span in mapas.HOJAS_POR_LIBRO[c]})
            out.append(OrderedDict([('es', span), ('en', ' / '.join(en))]))
    return out


def palabras(s):
    return len(re.findall(r'\w+', s))


def main():
    ap = argparse.ArgumentParser(description=__doc__.split('\n')[1])
    ap.add_argument('--origen-repo', default=REPO, help='raíz del repo de la que salen los dl/ ES')
    ap.add_argument('--salida', default=mapas.DATOS)
    ap.add_argument('--md', action='store_true', help='imprime la tabla §0 de F1-inventario-es.md')
    args = ap.parse_args()

    try:
        commit = subprocess.check_output(['git', '-C', REPO, 'rev-parse', '--short', 'HEAD'],
                                         stderr=subprocess.DEVNULL).decode().strip()
    except Exception:                                             # noqa: BLE001
        commit = None
    censo = OrderedDict()
    censo['meta'] = OrderedDict([
        ('generado', datetime.datetime.utcnow().replace(microsecond=0).isoformat() + 'Z'),
        ('script', 'scripts/productos-digitales/business-plans-en/extraer_textos.py'),
        ('origen', [mapas.PLANES[p]['dir_es'] for p in mapas.PLANES]),
        ('commit', commit),
        ('nota', 'Censo de los 4 xlsx v2.2 PUBLICADOS (plan financiero + checklist × FT/CAF). Los gates de la F2-EN '
                 'comparan contra este fichero, nunca contra recuentos fijos. formulas_texto = fórmulas con literales '
                 'de texto (las cubre mapas.FORMULAS_EN o mapas.CLAVES). tokens_docx = celdas de las cifras del docx.'),
    ])
    censo['libros'] = OrderedDict()
    todas = []
    for corto, (plan, f_es, f_en, tit) in mapas.LIBROS.items():
        p = os.path.join(args.origen_repo, mapas.PLANES[plan]['dir_es'], f_es)
        if not os.path.exists(p):
            sys.exit('falta ' + p)
        info, oc = censar_libro(p, f_es)
        censo['libros'][f_es] = info
        todas.extend(oc)
    censo['tokens_docx'] = resolver_tokens(censo)
    for info in censo['libros'].values():
        info.pop('_wsv')
    SUM = ('hojas', 'celdas_texto', 'celdas_formula', 'formulas_con_hoja', 'formulas_texto', 'dv', 'cf', 'merges',
           'hojas_protegidas', 'paneles', 'formatos_euro', 'verdes', 'desbloqueadas', 'tareas')
    censo['totales'] = OrderedDict((k, sum(i['totales'][k] for i in censo['libros'].values())) for k in SUM)

    # ---- cadenas únicas (las de zona celda a celda: una por celda)
    todas_hojas = mapas.todas_hojas_es()
    excl = OrderedDict()
    cad = OrderedDict()
    for texto, oc in todas:
        fijado = texto in mapas.FIJOS or (oc.get('f'), oc.get('h'), oc.get('c')) in mapas.POR_CELDA
        if oc['tr'] == 'localizar':
            cad.setdefault(('@', oc['f'], oc['h'], oc['c']), (texto, []))[1].append(oc)
            continue
        if (not tiene_letra(texto) and not RX_MONEDA.search(texto) and not fijado) or texto in INVARIABLES:
            excl[texto] = excl.get(texto, 0) + 1
            continue
        cad.setdefault(texto, (texto, []))[1].append(oc)
    cadenas = []
    for i, (clave, (texto, ocs)) in enumerate(cad.items(), 1):
        cortos = []
        for o in ocs:
            if o['f'] not in cortos:
                cortos.append(o['f'])
        planes = sorted({mapas.PLAN_DE[c] for c in cortos})
        trs = sorted({o['tr'] for o in ocs}, key=lambda t: PRIORIDAD[t])
        e = OrderedDict()
        e['id'] = 'c{:04d}'.format(i)
        e['es'] = texto
        e['grupo'] = None
        e['tratamiento'] = trs[0]
        if len(trs) > 1:
            e['tratamientos'] = trs
        if trs[0] in ('mapa-hoja',):
            e['en_mapa'] = ' / '.join(sorted({mapas.HOJAS_POR_LIBRO[c][texto] for c in cortos
                                              if texto in mapas.HOJAS_POR_LIBRO[c]}))
        elif trs[0] == 'mapa-clave':
            e['en_mapa'] = mapas.CLAVES[texto]
        elif trs[0] == 'regenerar' and texto in mapas.FIJOS:
            e['en_mapa'] = mapas.FIJOS[texto]
        e['planes'] = planes
        e['n_apariciones'] = len(ocs)
        e['palabras'] = palabras(texto)
        lim = [LIMITE[o['t']] for o in ocs if o['t'] in LIMITE]
        if lim:
            e['max_len'] = min(lim)
        ch = citas_hoja(texto, todas_hojas, cortos)
        if ch:
            e['cita_hojas'] = ch
        if trs[0] in ('traducir', 'localizar'):
            pi = pistas(texto)
            for o in ocs:
                extra = mapas.CELDA_PISTAS.get((o['f'], o['h'], o.get('c')))
                if extra and extra not in pi:
                    pi.append(extra)
                if o['tr'] == 'localizar':
                    pz = mapas.PISTA_D18 if o.get('det') == 'D18' else mapas.PISTA_D19
                    if pz not in pi:
                        pi.insert(0, pz)
            if pi:
                e['pistas'] = pi
        if e['tratamiento'] == 'localizar':
            e['por_celda'] = True
        e['sha1'] = hashlib.sha1(texto.encode('utf-8')).hexdigest()[:12]
        e['donde'] = ocs
        cadenas.append(e)

    for e in cadenas:
        if e['tratamiento'] in ('traducir', 'localizar'):
            e['grupo'] = mapas.grupo_de(e['planes'], e['es'], e['tratamiento'])
            r = mapas.REUSO.get(e['es']) if e['tratamiento'] == 'traducir' else None
            if r:                                          # ya traducida en otro conjunto (restpan: GM/GFT/GCAF)
                e['reuso'] = r
        else:
            e['grupo'] = 'GX'
    DESC = mapas.GRUPOS_DESC
    grupos = []
    for gid, desc in DESC.items():
        es = [e for e in cadenas if e['grupo'] == gid]
        g = OrderedDict([('id', gid), ('descripcion', desc), ('n_cadenas', len(es)),
                         ('n_palabras', sum(e['palabras'] for e in es))])
        for tipo in ('plan', 'check'):
            g['n_' + tipo] = sum(1 for e in es if any(mapas.TIPO_DE[o['f']] == tipo for o in e['donde']))
        g['n_localizar'] = sum(1 for e in es if e['tratamiento'] == 'localizar')
        g['n_max_len'] = sum(1 for e in es if 'max_len' in e)
        if gid == 'GX':
            por_tr = defaultdict(int)
            for e in es:
                por_tr[e['tratamiento']] += 1
            g['por_tratamiento'] = OrderedDict(sorted(por_tr.items()))
        grupos.append(g)
    orden_g = {g: k for k, g in enumerate(DESC)}
    cadenas.sort(key=lambda e: (orden_g[e['grupo']], int(e['id'][1:])))

    lit_formula = OrderedDict()
    for info in censo['libros'].values():
        for h, hd in info['hojas_detalle'].items():
            for lit, cs in hd['literales'].items():
                if mapas.necesita_clave(lit):
                    lit_formula.setdefault(lit, []).append('%s!%s!%s' % (info['corto'], h, ','.join(cs)))
    textos = OrderedDict()
    textos['meta'] = OrderedDict([
        ('generado', censo['meta']['generado']), ('script', censo['meta']['script']), ('commit', commit),
        ('libros', OrderedDict((i['corto'], OrderedDict([('es', f), ('en', i['nombre_en']), ('sha256', i['sha256'])]))
                               for f, i in censo['libros'].items())),
        ('n_cadenas', len(cadenas)),
        ('n_traducir', sum(1 for e in cadenas if e['tratamiento'] == 'traducir')),
        ('n_localizar', sum(1 for e in cadenas if e['tratamiento'] == 'localizar')),
        ('n_palabras_subagentes', sum(e['palabras'] for e in cadenas if e['grupo'] != 'GX')),
        ('n_apariciones', sum(e['n_apariciones'] for e in cadenas)),
        ('grupos', grupos),
        ('excluidas', OrderedDict([
            ('cadenas', excl),
            ('literales_formula', OrderedDict([
                ('n', len(lit_formula)),
                ('regla', 'NO van a los subagentes: los literales de texto de las fórmulas los cubre mapas.CLAVES (cortos) '
                          'o mapas.FORMULAS_EN (la fórmula entera, celda a celda, ya en EN)'),
                ('lista', lit_formula)])),
            ('regla', 'Fuera: cadenas sin letras (salvo zonas celda a celda) y las invariables %s.' % sorted(INVARIABLES)),
        ])),
        ('campos', OrderedDict([
            ('donde', 'f = libro (FTP/FTC/CAFP/CAFC), h = hoja ES, c = celda / sqref / campo; t = celda | hoja | dv-* | '
                      'cabecera-pie | docprops; tr = tratamiento de esa aparición; det = detalle'),
            ('en_mapa', 'valor EN obligatorio que dicta mapas.py (GX)'),
            ('cita_hojas', 'pestañas ES citadas entre comillas, con su EN'),
            ('pistas', 'decisiones de la SPEC que aplican a la cadena'),
            ('max_len', 'límite de Excel: 32 (título de DV) o 255 (mensaje de DV)'),
            ('por_celda', 'cadena de una tabla D18/D19: su EN vale SOLO para esa celda'),
        ])),
    ])
    textos['cadenas'] = cadenas

    os.makedirs(args.salida, exist_ok=True)
    for nombre, obj in (('textos_es.json', textos), ('censo_es.json', censo)):
        with open(os.path.join(args.salida, nombre), 'w', encoding='utf-8') as fh:
            json.dump(obj, fh, ensure_ascii=False, indent=1)
            fh.write('\n')
    print('textos_es.json: {} cadenas ({} traducir, {} localizar, {} palabras para subagentes) · {} apariciones · '
          '{} literales de fórmula fuera'.format(len(cadenas), textos['meta']['n_traducir'],
                                                textos['meta']['n_localizar'], textos['meta']['n_palabras_subagentes'],
                                                textos['meta']['n_apariciones'], len(lit_formula)))
    for g in grupos:
        print('  {:<4} {:>4} cadenas {:>6} palabras  plan {:>3} · check {:>3} · celda {:>3} · ≤len {:>3}'.format(
            g['id'], g['n_cadenas'], g['n_palabras'], g['n_plan'], g['n_check'], g['n_localizar'], g['n_max_len']))
    print('censo_es.json: ' + json.dumps(censo['totales'], ensure_ascii=False))
    for f, i in censo['libros'].items():
        print('  %-5s %s' % (i['corto'], json.dumps({k: v for k, v in i['totales'].items()
                                                    if not isinstance(v, (dict, list))}, ensure_ascii=False)))
    if args.md:
        print(tabla_md(censo))


COLS_MD = [('hojas', 'Hojas'), ('celdas_texto', 'Texto'), ('celdas_formula', 'Fórmulas'),
           ('formulas_con_hoja', 'Fx→hoja'), ('dv', 'DV'), ('cf', 'CF'), ('merges', 'Merges'), ('verdes', 'Verdes'),
           ('desbloqueadas', 'Desbloq.'), ('formatos_euro', 'Fmt €'), ('tareas', 'Tareas')]


def tabla_md(censo):
    out = ['| Libro | ' + ' | '.join(t for _, t in COLS_MD) + ' |', '|---|' + '---|' * len(COLS_MD)]
    for f, info in censo['libros'].items():
        out.append('| %s | ' % f[:-5] + ' | '.join(str(info['totales'][k]) for k, _ in COLS_MD) + ' |')
    return '\n'.join(out)


if __name__ == '__main__':
    main()
