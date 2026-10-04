#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
aplicar_en.py — Food Truck + Coffee Shop Business Plan Kit (EN) · montaje de los 4 xlsx EN desde los 4 xlsx ES
publicados (plan financiero + checklist de cada plan).

SPEC que manda: `SPEC.md` (§1 D6-D23, §1.1 E1/E2, §2-§4, §7 paso 3, §8). Sesión Claude Code, 4-oct-2026. Copia
adaptada de `financial-kit/aplicar_en.py` (misma maquinaria: renombrado de pestañas reescribiendo las referencias,
literales de fórmula/CF/DV por CLAVES, traducción POR APARICIÓN, formatos sin «€», US Letter + pie, docProps,
comillas rectas, huella de idempotencia e `inject_cache` AL FINAL en el mismo proceso). Aquí no se redacta nada: los
textos salen de `textos_en/{GM,GFT,GCAF}.json` (+ `overrides_*.json`, por celda) y de `mapas.py`.

    python3 aplicar_en.py                          # escribe en astro-site/public/dl/<slug>/ y cifras_caso.json
    python3 aplicar_en.py --dry-run [--salida DIR] [--idempotencia] [--json informe.json]
    python3 aplicar_en.py --dry-run --parcial      # deja en español lo que aún no tenga traducción (solo ensayo)

Los xlsx ES de `dl/plan-negocio-food-truck/` y `dl/plan-negocio-cafeteria/` se abren en SOLO LECTURA.

Orden por libro (F2-NOTAS §5.2), siempre desde una copia limpia del ES publicado (→ idempotente):
   1. E2 (mapas.PARCHES_FORMULA) en espacio ES;
   2. pestañas (mapas.HOJAS_POR_LIBRO) reescribiendo las referencias de fórmulas, DV y CF;
   3. mapas.FORMULAS_EN celda a celda (ya en EN); el resto de fórmulas, pestañas + CLAVES (un literal con letras sin
      clave aborta);
   4. textos POR APARICIÓN (textos_es.json `donde`): traducir/localizar → textos_en (o el override de ESA celda),
      mapa-clave → CLAVES, mapa-hoja → pestaña EN, regenerar → FIJOS / línea de versión / pie;
   5. DV (lista por CLAVES, referencias, mensajes por aparición) y CF (fórmula y `text` por CLAVES);
   6. VALORES_EN (D15-D16) + CALIBRACION_D15 (mapas) y el préstamo B31 = el que propone `Financing`, a 1,000;
   7. formatos D8 (sin «€»), US Letter en las 32 hojas;
   8. E1: el aviso D21 bajo la línea de versión (4 celdas);
   9. protección: viaja intacta con openpyxl (se comprueba en gates G1);
  10. docProps D20, comillas rectas, cobertura de «regenerar»; guardado con el nombre EN (D5);
  11. inject_cache AL FINAL sobre los 4, en el mismo proceso (nunca se vuelve a guardar con openpyxl después);
  12. calibración D15: lee la caché y exige DSCR mínimo ≥ 1.25, saldo mínimo > 0, cobertura ≥ 98 %, ningún sueldo
      bajo el suelo y holgura sobre el equilibrio de caja ≥ 15 % (los ratios del P&L se informan: pueden quedar
      ámbar con nota);
  13. cifras_caso.json con mapas.calcular_cifras() sobre la caché de los planes EN (lo lee ensamblar_docx.py).

El préstamo (B31 = CALIBRAR) se busca por punto fijo: se construyen los dos planes con una semilla, se recalcula,
se lee `Financing!B18` («loan that would match funding to the need»), se redondea a 1,000 por arriba y se repite
hasta que no cambia (2-3 vueltas; el fondo de maniobra incluye los intereses, por eso no basta una).

Térmica: todo en SERIE; `istats` antes de cada libro y de cada inject_cache (en el VPS no hay istats y no espera).
"""
import argparse
import copy
import datetime
import hashlib
import json
import logging
import math
import os
import re
import shutil
import subprocess
import sys
import tempfile
import time
from collections import OrderedDict

AQUI = os.path.dirname(os.path.abspath(__file__))
SCRIPTS = os.path.dirname(AQUI)
REPO = os.path.dirname(os.path.dirname(SCRIPTS))
sys.path.insert(0, AQUI)
sys.path.insert(1, SCRIPTS)

import openpyxl                                              # noqa: E402

import mapas                                                 # noqa: E402
import extraer_textos                                        # noqa: E402

logging.disable(logging.CRITICAL)

CLAVES = mapas.CLAVES
CAMPO_DV = {'dv-error-titulo': 'errorTitle', 'dv-error': 'error', 'dv-prompt-titulo': 'promptTitle', 'dv-prompt': 'prompt'}
INVARIABLES = set(extraer_textos.INVARIABLES)
RX_VERSION_ES = extraer_textos.RX_VERSION
VERDE = extraer_textos.VERDE
PLANES_CORTO = OrderedDict([('ft', 'FTP'), ('caf', 'CAFP')])
CHECK_CORTO = OrderedDict([('ft', 'FTC'), ('caf', 'CAFC')])
UMBRALES_D15 = OrderedDict([('dscr_min', 1.25), ('saldo_minimo', 0.0), ('cobertura_horas', 0.98),
                            ('holgura_caja_pct', 0.15)])


class Aborta(Exception):
    pass


# ==========================================================================
# Térmica (regla del Mac: se apaga por encima de ~65 °C; en el VPS no hay istats)
# ==========================================================================
def termica(umbral=62, reanuda=60):
    def leer():
        try:
            out = subprocess.run(['istats', 'cpu', 'temp', '--value-only'], stdout=subprocess.PIPE,
                                 stderr=subprocess.PIPE, universal_newlines=True, timeout=20).stdout
            return float(re.findall(r'[\d.]+', out)[0])
        except Exception:
            return None
    t = leer()
    if t is None or t < umbral:
        return t
    while t is not None and t >= reanuda:
        print('  [térmica] CPU %.1f °C: espero 30 s' % t, flush=True)
        time.sleep(30)
        t = leer()
    return t


# ==========================================================================
# Datos
# ==========================================================================
def destino(plan):
    return os.path.join(REPO, mapas.PLANES[plan]['dir_en'])


def cargar_datos(parcial=False):
    te = json.load(open(os.path.join(AQUI, 'textos_es.json'), encoding='utf-8'))
    censo = json.load(open(os.path.join(AQUI, 'censo_es.json'), encoding='utf-8'))
    por_id = {c['id']: c for c in te['cadenas']}
    tx = OrderedDict()
    for g in ('GM', 'GFT', 'GCAF'):
        p = os.path.join(AQUI, 'textos_en', g + '.json')
        if not os.path.exists(p):
            if parcial:
                continue
            raise Aborta('falta textos_en/%s.json' % g)
        for k, v in json.load(open(p, encoding='utf-8')).items():
            if k not in por_id:
                raise Aborta('textos_en/%s.json trae un id que no está en textos_es.json: %s' % (g, k))
            if por_id[k]['grupo'] != g:
                raise Aborta('%s es del grupo %s y viene en %s.json' % (k, por_id[k]['grupo'], g))
            tx[k] = v
    celdas, dvmsg = {}, {}
    for c in te['cadenas']:
        for d in c['donde']:
            if d['t'] == 'celda':
                celdas[(d['f'], d['h'], d['c'])] = (c['id'], d['tr'], c['es'])
            elif d['t'] in CAMPO_DV:
                dvmsg[(d['f'], d['h'], d['c'], CAMPO_DV[d['t']])] = (c['id'], d['tr'], c['es'])
    # overrides por libro: {id: {libro, hoja (ES), celda, texto}} → solo en ESA celda
    overrides = OrderedDict()
    for g in ('GFT', 'GCAF'):
        p = os.path.join(AQUI, 'textos_en', 'overrides_%s.json' % g)
        if not os.path.exists(p):
            continue
        for ident, o in json.load(open(p, encoding='utf-8')).items():
            k = (o['libro'], o['hoja'], o['celda'])
            if k not in celdas:
                raise Aborta('override %s: %s!%s (%s) no es una aparición de textos_es.json' % (ident, o['hoja'],
                                                                                              o['celda'], o['libro']))
            if celdas[k][0] != ident:
                raise Aborta('override %s: la celda %s!%s es de la cadena %s' % (ident, o['hoja'], o['celda'],
                                                                               celdas[k][0]))
            r = [x for x in mapas.restos_espanol(o['texto'])]
            if r or mapas.no_latinos(o['texto']):
                raise Aborta('override %s con restos de español: %r' % (ident, r[:2]))
            overrides[k] = (ident, o['texto'])
    faltan = sorted({c['id'] for c in te['cadenas'] if c['grupo'] in ('GM', 'GFT', 'GCAF') and c['id'] not in tx})
    if faltan and not parcial:
        raise Aborta('%d cadenas sin traducción (p. ej. %r)' % (len(faltan), faltan[:5]))
    return {'te': te, 'censo': censo, 'tx': tx, 'celdas': celdas, 'dvmsg': dvmsg, 'overrides': overrides,
            'faltan': faltan, 'parcial': parcial}


def censo_libro(datos, corto):
    for f, info in datos['censo']['libros'].items():
        if info['corto'] == corto:
            return info
    raise Aborta('libro %s fuera del censo' % corto)


# ==========================================================================
# Fórmulas: pestañas y literales
# ==========================================================================
RX_STR = re.compile(r'"(?:[^"]|"")*"')
RX_REF = re.compile(r"(?:'((?:[^']|'')+)'|([^\W\d][\w.]*))!")


def ref_hoja(nombre):
    """Nombre de hoja tal como va en una fórmula (entre comillas si hace falta)."""
    if re.fullmatch(r'[A-Za-z_][A-Za-z0-9_.]*', nombre):
        return nombre
    return "'" + nombre.replace("'", "''") + "'"


def clave_en(lit, donde):
    """Literal ES → EN por CLAVES: estricto con los que llevan letras o «€»; los demás no cambian."""
    if not mapas.necesita_clave(lit):
        return lit
    if lit not in CLAVES:
        raise Aborta('literal sin clave EN %r en %r' % (lit, donde[:90]))
    return CLAVES[lit]


class Transformador:
    """Reescribe una fórmula ES a EN: referencias de hoja (fuera de literales) y literales (dentro)."""

    def __init__(self, mapa_hojas):
        self.mapa = mapa_hojas
        self.en = set(mapa_hojas.values())
        self.hojas_citadas = 0
        self.literales = 0

    def refs(self, s):
        def rep(m):
            nombre = (m.group(1) if m.group(1) is not None else m.group(2)).replace("''", "'")
            if nombre in self.mapa:
                self.hojas_citadas += 1
                return ref_hoja(self.mapa[nombre]) + '!'
            if nombre in self.en:
                return m.group(0)
            raise Aborta('hoja desconocida en fórmula: %r en %r' % (nombre, s[:80]))
        return RX_REF.sub(rep, s)

    def formula(self, f):
        out, pos = [], 0
        for m in RX_STR.finditer(f):
            out.append(self.refs(f[pos:m.start()]))
            lit = m.group(0)[1:-1].replace('""', '"')
            nuevo = clave_en(lit, f)
            if nuevo != lit:
                self.literales += 1
            out.append('"' + nuevo.replace('"', '""') + '"')
            pos = m.end()
        out.append(self.refs(f[pos:]))
        return ''.join(out)

    def celda(self, v):
        return '=' + self.formula(v[1:]) if v.startswith('=') else self.formula(v)


def transformar_es(corto, hoja_es, coord, f_es):
    """Fórmula EN ESPERADA de una celda ES (la usan aplicar_en y gates_en): FORMULAS_EN si la cubre; si no, E2 en
    espacio ES y luego pestañas + CLAVES."""
    cub = _CUBIERTAS.get((corto, hoja_es, coord))
    if cub is not None:
        return cub
    f = mapas.parchear_formula_es(corto, hoja_es, coord, f_es)
    return Transformador(mapas.HOJAS_POR_LIBRO[corto]).celda(f)


_CUBIERTAS = mapas.celdas_formulas_en()


def rango(ref):
    c1, r1, c2, r2 = openpyxl.utils.cell.range_boundaries(ref)
    for r in range(r1, r2 + 1):
        for c in range(c1, c2 + 1):
            yield '%s%d' % (openpyxl.utils.get_column_letter(c), r)


def renombrar(wb, corto):
    mapa = mapas.HOJAS_POR_LIBRO[corto]
    if wb.sheetnames != mapas.HOJAS_ES_POR_LIBRO[corto]:
        raise Aborta('%s: pestañas ES %r ≠ censo %r' % (corto, wb.sheetnames, mapas.HOJAS_ES_POR_LIBRO[corto]))
    for ws in wb.worksheets:
        ws.title = mapa[ws.title]
    return mapa


def reescribir_formulas(L):
    """E2 en espacio ES, FORMULAS_EN celda a celda y, en el resto, pestañas + CLAVES."""
    n = n_fx = n_e2 = 0
    aplicadas = set()
    for ws in L.wb.worksheets:
        h_es = L.inv[ws.title]
        for c in list(ws._cells.values()):
            if c.data_type != 'f' or not isinstance(c.value, str):
                continue
            k = (L.corto, h_es, c.coordinate)
            if k in _CUBIERTAS:
                if c.coordinate not in L.cen['hojas_detalle'][h_es]['formulas_texto']:
                    raise Aborta('%s %s!%s: FORMULAS_EN sobre una celda que no es fórmula de texto' % k)
                c.value = _CUBIERTAS[k]
                aplicadas.add(k)
                n_fx += 1
                n += 1
                continue
            try:
                v = mapas.parchear_formula_es(L.corto, h_es, c.coordinate, c.value)
            except ValueError as e:
                raise Aborta(str(e))
            if v != c.value:
                n_e2 += 1
            nuevo = L.T.celda(v)
            if nuevo != c.value:
                c.value = nuevo
                n += 1
    esperadas = {k for k in _CUBIERTAS if k[0] == L.corto}
    if aplicadas != esperadas:
        raise Aborta('%s: FORMULAS_EN sin aplicar: %r' % (L.corto, sorted(esperadas - aplicadas)[:5]))
    e2 = sum(1 for k in mapas.PARCHES_FORMULA if k[0] == L.corto)
    if n_e2 != e2:
        raise Aborta('%s: E2 aplicado en %d celdas y se esperaban %d' % (L.corto, n_e2, e2))
    L.rep.update(formulas_reescritas=n, formulas_en=n_fx, e2=n_e2)


# ==========================================================================
# Un libro en construcción
# ==========================================================================
class Libro:
    def __init__(self, wb, corto, datos, rep):
        self.wb, self.corto, self.datos, self.rep = wb, corto, datos, rep
        self.plan = mapas.PLAN_DE[corto]
        self.cen = censo_libro(datos, corto)
        self.mapa = renombrar(wb, corto)
        self.inv = {en: es for es, en in self.mapa.items()}
        self.T = Transformador(self.mapa)
        self.escritas = set()
        self.faltan = rep.setdefault('faltan', [])
        self.overrides_usados = set()

    def ws(self, hoja_es):
        return self.wb[self.mapa[hoja_es]]

    def poner(self, hoja_es, coord, valor, es_esperado=None):
        c = self.ws(hoja_es)[coord]
        if c.data_type == 'f' or (isinstance(c.value, str) and c.value.startswith('=')):
            raise Aborta('%s %s!%s es una fórmula (%r): una capa no la pisa' % (self.corto, hoja_es, coord, c.value))
        if es_esperado is not None and not es_esperado(c.value):
            raise Aborta('%s %s!%s: el ES trae %r, no lo esperado' % (self.corto, hoja_es, coord, c.value))
        c.value = valor
        self.escritas.add((hoja_es, coord))
        return c

    def en_de(self, ident, tr, es, lugar, celda=None):
        if celda is not None and celda in self.datos['overrides']:
            self.overrides_usados.add(celda)
            return self.datos['overrides'][celda][1]
        if tr in ('traducir', 'localizar'):
            en = self.datos['tx'].get(ident)
            if en is None:
                self.faltan.append('%s %s %r' % (lugar, ident, es[:60]))
                return es
            return en
        if tr == 'mapa-clave':
            return CLAVES[es]
        if tr == 'mapa-hoja':
            return mapas.HOJAS_POR_LIBRO[self.corto][es]
        return None                                              # regenerar: lo escribe una capa


def traducir_textos(L, mes):
    n = 0
    regenerar = []
    for ws in L.wb.worksheets:
        h_es = L.inv[ws.title]
        for c in list(ws._cells.values()):
            v = c.value
            if not isinstance(v, str) or c.data_type == 'f':
                continue
            if RX_VERSION_ES.match(v):
                L.poner(h_es, c.coordinate, mapas.version_en(L.plan, mes))
                n += 1
                continue
            k = (L.corto, h_es, c.coordinate)
            if k not in L.datos['celdas']:
                if (any(ch.isalpha() for ch in v) or '€' in v) and v not in INVARIABLES:
                    raise Aborta('%s %s!%s: texto %r fuera de textos_es.json' % (L.corto, h_es, c.coordinate, v[:60]))
                continue                                         # sin letras / invariable (☐, ✓, —, OK…)
            ident, tr, es = L.datos['celdas'][k]
            if es != v:
                raise Aborta('%s %s!%s: textos_es.json dice %r y el ES trae %r' % (L.corto, h_es, c.coordinate,
                                                                                es[:40], v[:40]))
            en = L.en_de(ident, tr, es, '%s %s!%s' % (L.corto, ws.title, c.coordinate), celda=k)
            if en is None:
                regenerar.append((h_es, c.coordinate))
                continue
            if en != v:
                c.value = en
                n += 1
        for parte in ('oddHeader', 'oddFooter', 'evenHeader', 'evenFooter', 'firstHeader', 'firstFooter'):
            hf = getattr(ws, parte)
            for pos in ('left', 'center', 'right'):
                t = getattr(hf, pos).text
                if not t:
                    continue
                if t != mapas.FOOTER_ES:
                    raise Aborta('%s %s %s.%s inesperado: %r' % (L.corto, ws.title, parte, pos, t))
                getattr(hf, pos).text = mapas.FOOTER_EN
    sin_usar = [k for k in L.datos['overrides'] if k[0] == L.corto and k not in L.overrides_usados]
    if sin_usar:
        raise Aborta('%s: overrides sin aplicar: %r' % (L.corto, sin_usar))
    L.rep['textos_traducidos'] = n
    L.rep['overrides'] = len(L.overrides_usados)
    return regenerar


def capa_fijos(L, regenerar):
    """FIJOS: cadenas fijadas por la SPEC, por TEXTO, en toda aparición «regenerar» aún sin escribir."""
    n = 0
    for h_es, coord in regenerar:
        if (h_es, coord) in L.escritas:
            continue
        v = L.ws(h_es)[coord].value
        if v in mapas.FIJOS:
            L.poner(h_es, coord, mapas.FIJOS[v], lambda x: isinstance(x, str))
            n += 1
    L.rep['fijos'] = n


def reescribir_dv_cf(L):
    n = OrderedDict([('dv_listas', 0), ('dv_personalizadas', 0), ('dv_ref', 0), ('dv_mensajes', 0), ('cf_reglas', 0),
                     ('cf_text', 0)])
    for ws in L.wb.worksheets:
        h_es = L.inv[ws.title]
        sqs_censo = {d['sqref'] for d in L.cen['hojas_detalle'][h_es]['dv']}
        for dv in ws.data_validations.dataValidation:
            sq = str(dv.sqref)
            if sq not in sqs_censo:
                raise Aborta('%s %s: DV %s fuera del censo' % (L.corto, h_es, sq))
            f1 = dv.formula1 or ''
            if dv.type == 'list' and f1.startswith('"'):
                items = [clave_en(x, 'DV %s %s' % (h_es, sq)) if x else x for x in f1.strip('"').split(',')]
                dv.formula1 = '"' + ','.join(items) + '"'
                if len(dv.formula1) > 257:
                    raise Aborta('%s %s: DV %s > 255 caracteres' % (L.corto, h_es, sq))
                n['dv_listas'] += 1
            elif dv.type == 'custom' and f1:
                dv.formula1 = L.T.formula(f1)
                n['dv_personalizadas'] += 1
            elif f1:
                dv.formula1 = L.T.formula(f1)
                if dv.formula2:
                    dv.formula2 = L.T.formula(dv.formula2)
                n['dv_ref'] += 1
            for attr in ('error', 'errorTitle', 'prompt', 'promptTitle'):
                v = getattr(dv, attr)
                if not v:
                    continue
                k = (L.corto, h_es, sq, attr)
                if k not in L.datos['dvmsg']:
                    raise Aborta('%s %s DV %s %s fuera de textos_es.json: %r' % (L.corto, h_es, sq, attr, v[:50]))
                ident, tr, es = L.datos['dvmsg'][k]
                if es != v:
                    raise Aborta('%s %s DV %s %s: textos_es.json dice %r' % (L.corto, h_es, sq, attr, es[:40]))
                en = L.en_de(ident, tr, es, '%s DV %s %s' % (ws.title, sq, attr))
                if en is None:
                    raise Aborta('%s %s DV %s %s marcado «regenerar»' % (L.corto, h_es, sq, attr))
                setattr(dv, attr, en)
                n['dv_mensajes'] += 1
        for cf in ws.conditional_formatting:
            for regla in cf.rules:
                if regla.formula:
                    regla.formula = [L.T.formula(x) for x in regla.formula]
                    n['cf_reglas'] += 1
                if regla.text:
                    regla.text = clave_en(regla.text, 'CF %s %s' % (ws.title, cf.sqref))
                    n['cf_text'] += 1
    L.rep.update(n)


def valores_libro(corto, prestamos):
    """VALORES_EN (research §7) + CALIBRACION_D15, con el préstamo resuelto. (hoja ES, celda) → valor."""
    out = OrderedDict()
    for (c, h, ref), (v, _rot) in mapas.VALORES_EN.items():
        if c != corto:
            continue
        out[(h, ref)] = v
    for (c, h, ref), (v, _motivo) in mapas.CALIBRACION_D15.items():
        if c != corto:
            continue
        if (h, ref) not in out:
            raise Aborta('CALIBRACION_D15 %s %s!%s no está en VALORES_EN' % (c, h, ref))
        out[(h, ref)] = v
    for k, v in out.items():
        if v == mapas.CALIBRAR:
            out[k] = prestamos[mapas.PLAN_DE[corto]]
    return out


def capa_valores(L, prestamos):
    n = 0
    num = lambda v: v is None or (isinstance(v, (int, float)) and not isinstance(v, bool))   # noqa: E731
    for (h, ref), v in valores_libro(L.corto, prestamos).items():
        L.poner(h, ref, v, num)
        n += 1
    L.rep['valores'] = n


def capa_textos_nuevos(L):
    """E1 (D21): el aviso legal en la fila de debajo de la línea de versión, con su mismo estilo."""
    n = 0
    for (c, h, coord), (texto, estilo) in mapas.TEXTOS_NUEVOS.items():
        if c != L.corto:
            continue
        ws = L.ws(h)
        cel = ws[coord]
        if cel.value is not None:
            raise Aborta('%s %s!%s no está vacía en el ES (%r): E1 no la pisa' % (c, h, coord, cel.value))
        ver = ws[estilo].value
        if not (isinstance(ver, str) and ver.startswith('Version ')):
            raise Aborta('%s: la línea de versión no está en %s (%r)' % (c, estilo, ver))
        cel._style = copy.copy(ws[estilo]._style)
        al = copy.copy(cel.alignment)                # el aviso se ajusta dentro de su celda: no desborda al imprimir
        al.wrap_text = True
        cel.alignment = al
        cel.value = texto
        alto = round(lineas_estimadas(texto, ws, cel) * pt_de(cel) * 1.36 * 4) / 4.0
        ws.row_dimensions[cel.row].height = max(alto, ws.row_dimensions[cel.row].height or 0)
        L.escritas.add((h, coord))
        n += 1
    L.rep['textos_nuevos'] = n


def pt_de(cel):
    return cel.font.sz if cel.font is not None and cel.font.sz else 11


def lineas_estimadas(texto, ws, cel):
    """Líneas que ocupa un texto ajustado en su columna (estimación: un carácter por unidad de ancho a 11 pt, con un
    15 % de holgura porque el texto real es más estrecho que el «0» de referencia)."""
    ancho = ws.column_dimensions[cel.column_letter].width or 8.43
    cpl = max(1, int(ancho * 11.0 / pt_de(cel) * 1.15))
    return sum(max(1, -(-len(seg) // cpl)) for seg in texto.split('\n'))


def capa_altos(L, wes):
    """Filas de texto ajustado cuyo EN ocupa más líneas que el ES y que su alto: se suben (nunca se bajan). Las
    celdas combinadas no se miden. Sin esto, una nota EN más larga que la ES se corta en pantalla e impresa."""
    n = 0
    for ws in L.wb.worksheets:
        a = wes[L.inv[ws.title]]
        comb = {c for m in ws.merged_cells.ranges for fila in ws[m.coord] for c in (x.coordinate for x in fila)}
        filas = {}
        for c in ws._cells.values():
            if not isinstance(c.value, str) or c.data_type == 'f' or not (c.alignment and c.alignment.wrap_text):
                continue
            if c.coordinate in comb:
                continue
            ve = a[c.coordinate].value
            ne = lineas_estimadas(c.value, ws, c)
            ns = lineas_estimadas(ve, ws, c) if isinstance(ve, str) else 0
            f = filas.setdefault(c.row, [0, 0, 11])
            f[0], f[1], f[2] = max(f[0], ne), max(f[1], ns), max(f[2], pt_de(c))
        for r, (ne, ns, pt) in sorted(filas.items()):
            h = ws.row_dimensions[r].height or 15
            if ne > max(ns, int(h / (pt * 1.25))):
                ws.row_dimensions[r].height = max(h, round(ne * pt * 1.36 * 4) / 4.0)
                n += 1
    L.rep['altos_subidos'] = n


def formatos_y_papel(L):
    n = OrderedDict((k, 0) for k in mapas.FORMATOS_EN)
    for ws in L.wb.worksheets:
        for c in list(ws._cells.values()):
            f = c.number_format
            if f in mapas.FORMATOS_EN:
                c.number_format = mapas.FORMATOS_EN[f]
                n[f] += 1
            elif '€' in (f or ''):
                raise Aborta('%s %s!%s: formato con € sin mapa: %r' % (L.corto, ws.title, c.coordinate, f))
        ws.page_setup.paperSize = mapas.PAPEL_EN
    # formatos propios que ya no usa ninguna celda (los ES con €) → su variante sin €
    nf = L.wb._number_formats
    usados = set(nf)
    for i, f in enumerate(list(nf)):
        if '€' in f:
            nuevo = mapas.FORMATOS_EN.get(f, '#,##0.00')
            while nuevo in usados:
                nuevo += ';@'
            usados.add(nuevo)
            list.__setitem__(nf, i, nuevo)
    if len(set(nf)) != len(nf):
        raise Aborta('formatos propios repetidos: %r' % list(nf))
    nf._dict = {v: i for i, v in enumerate(nf)}
    nf.clean = True
    L.rep['formatos'] = OrderedDict((k, v) for k, v in n.items() if v)


COMILLAS = {'\u2019': "'", '\u2018': "'", '\u201c': '"', '\u201d': '"'}


def norm_comillas(v):
    for a, b in COMILLAS.items():
        v = v.replace(a, b)
    return v


def normalizar_comillas(L):
    n = 0
    for ws in L.wb.worksheets:
        for c in list(ws._cells.values()):
            if isinstance(c.value, str) and c.data_type != 'f' and any(a in c.value for a in COMILLAS):
                c.value = norm_comillas(c.value)
                n += 1
        for dv in ws.data_validations.dataValidation:
            for attr in ('error', 'errorTitle', 'prompt', 'promptTitle'):
                v = getattr(dv, attr)
                if v and any(a in v for a in COMILLAS):
                    setattr(dv, attr, norm_comillas(v))
                    n += 1
    L.rep['comillas_normalizadas'] = n


def metadatos(L):
    p = L.wb.properties
    for k, v in mapas.docprops_en(L.corto).items():
        setattr(p, k, v)
    L.rep['docprops'] = mapas.docprops_en(L.corto)


def construir_libro(corto, carpeta, datos, mes, prestamos):
    plan = mapas.PLAN_DE[corto]
    src = mapas.ruta_es(corto, REPO)
    dst = os.path.join(carpeta, mapas.LIBROS[corto][2])
    rep = OrderedDict([('corto', corto), ('es', os.path.relpath(src, REPO)), ('en', mapas.LIBROS[corto][2])])
    wb = openpyxl.load_workbook(src)
    L = Libro(wb, corto, datos, rep)                                            # 2 pestañas
    reescribir_formulas(L)                                                      # 1, 3
    regenerar = traducir_textos(L, mes)                                         # 4
    capa_fijos(L, regenerar)
    reescribir_dv_cf(L)                                                         # 5
    if mapas.TIPO_DE[corto] == 'plan':
        capa_valores(L, prestamos)                                              # 6
    formatos_y_papel(L)                                                         # 7
    capa_textos_nuevos(L)                                                       # 8
    capa_altos(L, openpyxl.load_workbook(src))                                  # alto de las filas que el EN alarga
    metadatos(L)                                                                # 10
    normalizar_comillas(L)
    sin_capa = [x for x in regenerar if x not in L.escritas]
    if sin_capa:
        raise Aborta('%s: apariciones «regenerar» que ninguna capa escribió: %r' % (corto, sin_capa[:8]))
    rep['regeneradas'] = len(regenerar)
    rep['refs_hoja'] = L.T.hojas_citadas
    rep['literales'] = L.T.literales
    if rep['faltan'] and not datos['parcial']:
        raise Aborta('%s: %d textos sin traducción:\n  %s' % (corto, len(rep['faltan']), '\n  '.join(rep['faltan'][:20])))
    os.makedirs(os.path.dirname(dst), exist_ok=True)
    wb.save(dst)
    return dst, rep


def construir(carpetas, datos, mes, prestamos, cortos=None, log=print):
    informe = []
    for corto in (cortos or list(mapas.LIBROS)):
        termica()
        dst, rep = construir_libro(corto, carpetas[mapas.PLAN_DE[corto]], datos, mes, prestamos)
        informe.append(rep)
        log('  %-42s fórmulas %4d (EN %2d, E2 %d) · refs %4d · lit %3d · textos %4d · fijos %3d · valores %3d · '
            'faltan %d' % (rep['en'], rep['formulas_reescritas'], rep['formulas_en'], rep['e2'], rep['refs_hoja'],
                           rep['literales'], rep['textos_traducidos'], rep['fijos'], rep.get('valores', 0),
                           len(rep['faltan'])))
    return informe


# ==========================================================================
# inject_cache y lectura de la caché
# ==========================================================================
def inyectar(path):
    import inject_cache as IC
    return IC.inject(path)


def inject_cache(rutas, log=print):
    fuera = []
    for p in rutas:
        termica()
        try:
            nf, inj, fail = inyectar(p)
            linea = '%s: formulas=%d cache_inyectado=%d fallos_pycel=%d' % (os.path.basename(p), nf, inj, fail)
            ok = True
        except Exception as e:                                # noqa: BLE001
            linea, ok, fail = '%s: EXCEPCIÓN %r' % (os.path.basename(p), e), False, -1
        log('    ' + linea)
        fuera.append({'fichero': os.path.basename(p), 'salida': linea, 'ok': ok, 'fallos_pycel': fail})
    return fuera


class Cache:
    """Lector de la caché de un plan EN por (hoja ES, celda)."""

    def __init__(self, path, corto):
        self.corto = corto
        self.wv = openpyxl.load_workbook(path, data_only=True)

    def __call__(self, hoja_es, celda):
        return self.wv[mapas.hoja_en(self.corto, hoja_es)][celda].value


def fila_de(censo_libro_, hoja_es, rotulo):
    for fila, et in censo_libro_['hojas_detalle'][hoja_es]['etiquetas'].items():
        if (et or '').startswith(rotulo):
            return int(fila)
    raise Aborta('rótulo %r no está en %s' % (rotulo, hoja_es))


def prestamo_propuesto(leer, cen):
    """Financing!B18 (préstamo que ajustaría el origen a la necesidad) → redondeado a 1,000 por arriba."""
    f = fila_de(cen, 'Financiación', 'Préstamo que ajustaría el origen a la necesidad')
    v = leer('Financiación', 'B%d' % f)
    if not isinstance(v, (int, float)):
        raise Aborta('Financing!B%d sin valor en la caché: %r' % (f, v))
    return int(math.ceil(round(v, 6) / 1000.0) * 1000)


def d15(leer, plan, datos):
    """Métricas D15 del caso base y ratios del P&L (año 1). Devuelve (métricas, fallos)."""
    corto = PLANES_CORTO[plan]
    cen = censo_libro(datos, corto)
    tok = datos['censo']['tokens_docx'][plan]
    m = OrderedDict()
    for t in UMBRALES_D15:
        m[t] = leer(tok[t]['hoja_es'], tok[t]['celda'])
    fallos = []
    for t, u in UMBRALES_D15.items():
        v = m[t]
        if not isinstance(v, (int, float)):
            fallos.append('%s sin valor numérico (%r)' % (t, v))
        elif (v <= u if t == 'saldo_minimo' else v < u - 1e-9):
            fallos.append('%s = %.4f %s %.4f' % (t, v, '≤' if t == 'saldo_minimo' else '<', u))
    # sueldos frente al suelo (CF de Personal!D: D/B × pagas < MAX(B22, B68) × C)
    f_tot = fila_de(cen, 'Personal', 'TOTAL PLANTILLA')
    pagas = leer('0. Supuestos', 'B21')
    suelo = max(leer('0. Supuestos', 'B22') or 0, leer('0. Supuestos', 'B68') or 0)
    bajo = []
    for f in range(5, f_tot):
        b, c, d = (leer('Personal', '%s%d' % (col, f)) for col in 'BCD')
        if not isinstance(d, (int, float)):
            continue
        if (d / (b or 1)) * pagas < suelo * c - 1e-6:
            bajo.append('Staffing fila %d: %.0f/persona × %d < %.0f × %.2f' % (f, d / (b or 1), pagas, suelo, c))
    m['sueldos_bajo_suelo'] = bajo
    fallos += bajo
    # ratios del P&L (informativos: pueden quedar ámbar con nota)
    rat = OrderedDict()
    for rot, sentido in (('Margen bruto / Ventas', '>='), ('Coste de mercancía / Ventas', '<='),
                         ('Coste de personal / Ventas', '<='), ('Aparcamiento / Ventas', '<='),
                         ('Alquiler / Ventas', '<='), ('Resultado neto / Ventas', '>=')):
        try:
            f = fila_de(cen, 'PyG 3 Años', rot)
        except Aborta:
            continue
        vals = [leer('PyG 3 Años', '%s%d' % (col, f)) for col in 'BCD']
        umbral = leer('PyG 3 Años', 'E%d' % f)
        ok = [(v >= umbral) if sentido == '>=' else (v <= umbral) for v in vals]
        rat[rot] = OrderedDict([('valores', vals), ('umbral', umbral), ('sentido', sentido), ('cumple', ok)])
    m['ratios_pyg'] = rat
    for t in ('ventas_a1', 'inversion_total', 'necesidad_caja', 'prestamo', 'fondos_propios', 'clientes_dia',
              'margen_bruto_pct', 'resultado_neto_a1', 'margen_neto_a1', 'payback', 'dscr_a1', 'equilibrio_caja_clientes_dia'):
        m[t] = leer(tok[t]['hoja_es'], tok[t]['celda'])
    return m, fallos


def buscar_prestamos(datos, mes, semilla, log=print):
    """Punto fijo del préstamo de los dos planes (B31 = CALIBRAR)."""
    prestamos = dict(semilla)
    tmp = tempfile.mkdtemp(prefix='bp-prestamo-')
    try:
        for vuelta in range(1, 7):
            carpetas = {p: tmp for p in mapas.PLANES}
            construir(carpetas, datos, mes, prestamos, cortos=list(PLANES_CORTO.values()), log=lambda *_: None)
            rutas = [os.path.join(tmp, mapas.LIBROS[c][2]) for c in PLANES_CORTO.values()]
            inject_cache(rutas, log=lambda *_: None)
            nuevo = {}
            for plan, corto in PLANES_CORTO.items():
                leer = Cache(os.path.join(tmp, mapas.LIBROS[corto][2]), corto)
                nuevo[plan] = prestamo_propuesto(leer, censo_libro(datos, corto))
            log('  vuelta %d: préstamo %s → propuesto %s' % (vuelta, prestamos, nuevo))
            if nuevo == prestamos:
                return prestamos
            prestamos = nuevo
        raise Aborta('el préstamo no converge en 6 vueltas: %r' % prestamos)
    finally:
        shutil.rmtree(tmp, ignore_errors=True)


def semilla_prestamos():
    p = os.path.join(AQUI, 'cifras_caso.json')
    try:
        c = json.load(open(p, encoding='utf-8'))
        return {plan: int(c[plan]['prestamo']['valor']) for plan in mapas.PLANES}
    except Exception:
        return {'ft': 80000, 'caf': 140000}


def escribir_cifras(carpetas, datos, destino_json):
    out = OrderedDict([('meta', OrderedDict([
        ('generado', datetime.datetime.utcnow().replace(microsecond=0).isoformat() + 'Z'),
        ('script', 'scripts/productos-digitales/business-plans-en/aplicar_en.py'),
        ('fuentes', OrderedDict((plan, mapas.LIBROS[corto][2]) for plan, corto in PLANES_CORTO.items())),
        ('nota', 'Cifras del caso de ejemplo US leídas de la caché de los xlsx EN calibrados (D15, D22). '
                 'Las lee ensamblar_docx.py; las fichas de landing EN citan las mismas.')]))])
    for plan, corto in PLANES_CORTO.items():
        leer = Cache(os.path.join(carpetas[plan], mapas.LIBROS[corto][2]), corto)
        out[plan] = mapas.calcular_cifras(plan, leer, datos['censo']['tokens_docx'][plan])
    with open(destino_json, 'w', encoding='utf-8') as fh:
        json.dump(out, fh, ensure_ascii=False, indent=1, default=float)
        fh.write('\n')
    return out


# ==========================================================================
# Idempotencia (huella comparable, sin fechas de guardado)
# ==========================================================================
def digest(path):
    wb = openpyxl.load_workbook(path)
    p = wb.properties
    fuera = {'_props': (p.title, p.subject, p.keywords, p.description, p.category, p.creator)}
    for ws in wb.worksheets:
        celdas = {}
        for c in ws._cells.values():
            if c.value is None and not (c.fill is not None and c.fill.fill_type == 'solid'):
                continue
            relleno = str(c.fill.fgColor.rgb) if c.fill is not None and c.fill.fill_type == 'solid' else None
            celdas[c.coordinate] = (repr(c.value), c.number_format, relleno, bool(c.protection.locked),
                                    bool(c.alignment.wrap_text))
        fuera[ws.title] = {
            'celdas': celdas,
            'merges': sorted(str(m) for m in ws.merged_cells.ranges),
            'dv': sorted('%s:%s:%s:%s:%s:%s:%s' % (dv.type, dv.formula1, dv.formula2, dv.sqref, dv.error,
                                                   dv.errorTitle, dv.prompt)
                         for dv in ws.data_validations.dataValidation),
            'cf': sorted('%s:%s:%s' % (cf.sqref, [r.formula for r in cf.rules], [r.text for r in cf.rules])
                         for cf in ws.conditional_formatting),
            'prot': (bool(ws.protection.sheet), ws.protection.password), 'area': str(ws.print_area),
            'papel': ws.page_setup.paperSize, 'pie': ws.oddFooter.center.text,
            'alto': sorted((k, d.height) for k, d in ws.row_dimensions.items() if d.height),
            'ancho': sorted((k, d.width) for k, d in ws.column_dimensions.items() if d.width),
        }
    return fuera


def diferencias(a, b, fichero):
    out = []
    if a['_props'] != b['_props']:
        out.append('%s: _props' % fichero)
    for hoja in sorted((set(a) | set(b)) - {'_props'}):
        if hoja not in a or hoja not in b:
            out.append('%s:%s: hoja solo en una pasada' % (fichero, hoja))
            continue
        for k in ('merges', 'dv', 'cf', 'prot', 'area', 'papel', 'pie', 'alto', 'ancho'):
            if a[hoja][k] != b[hoja][k]:
                out.append('%s:%s: cambia %s' % (fichero, hoja, k))
        ca, cb = a[hoja]['celdas'], b[hoja]['celdas']
        for coord in sorted(set(ca) | set(cb)):
            if ca.get(coord) != cb.get(coord):
                out.append('%s:%s!%s: %s → %s' % (fichero, hoja, coord, ca.get(coord), cb.get(coord)))
    return out


def sha256(p):
    h = hashlib.sha256()
    with open(p, 'rb') as fh:
        for b in iter(lambda: fh.read(1 << 16), b''):
            h.update(b)
    return h.hexdigest()


def main():
    ap = argparse.ArgumentParser(description='Food Truck + Coffee Shop Business Plan Kit (EN): montaje de los 4 xlsx')
    ap.add_argument('--dry-run', action='store_true', help='escribe en el scratchpad (o --salida), no en dl/')
    ap.add_argument('--salida', default=None, help='carpeta del dry-run (subcarpetas <slug>/)')
    ap.add_argument('--mes', default=mapas.MES_VERSION.split()[0], help='mes de la línea de versión D20')
    ap.add_argument('--json', default=None, help='informe JSON')
    ap.add_argument('--idempotencia', action='store_true', help='2.ª construcción en un clon y 0 diferencias')
    ap.add_argument('--parcial', action='store_true', help='(solo --dry-run) deja en ES lo que no tenga traducción')
    ap.add_argument('--prestamo-ft', type=int, default=None, help='(depuración) préstamo fijo, sin punto fijo')
    ap.add_argument('--prestamo-caf', type=int, default=None)
    args = ap.parse_args()
    if args.parcial and not args.dry_run:
        raise SystemExit('--parcial solo vale con --dry-run')
    mes = '%s %s' % (args.mes, mapas.MES_VERSION.split()[1])
    if args.dry_run:
        base = os.path.abspath(args.salida or os.path.join(tempfile.gettempdir(), 'bp-dryrun'))
        carpetas = {p: os.path.join(base, mapas.PLANES[p]['slug']) for p in mapas.PLANES}
    else:
        carpetas = {p: destino(p) for p in mapas.PLANES}
    for p, c in carpetas.items():
        if os.path.abspath(c) == os.path.abspath(os.path.join(REPO, mapas.PLANES[p]['dir_es'])):
            raise SystemExit('ABORTADO: la salida no puede ser la carpeta ES')

    t0 = time.time()
    try:
        datos = cargar_datos(args.parcial)
        if args.prestamo_ft and args.prestamo_caf:
            prestamos = {'ft': args.prestamo_ft, 'caf': args.prestamo_caf}
            print('== 0/4 · préstamo fijado a mano: %s' % prestamos, flush=True)
        else:
            print('== 0/4 · préstamo por punto fijo (B31 = Financing!B18 a 1,000 por arriba)', flush=True)
            prestamos = buscar_prestamos(datos, mes, semilla_prestamos())
        print('== 1/4 · montaje de los 4 libros → %s (versión «%s»)' % (sorted(set(carpetas.values())), mes), flush=True)
        informe = construir(carpetas, datos, mes, prestamos)
    except Aborta as e:
        raise SystemExit('ABORTADO: %s' % e)

    rutas = [os.path.join(carpetas[mapas.PLAN_DE[c]], mapas.LIBROS[c][2]) for c in mapas.LIBROS]
    idem = {'ejecutada': False}
    if args.idempotencia:
        print('== 2/4 · idempotencia: 2.ª construcción en un clon', flush=True)
        clon = tempfile.mkdtemp(prefix='bp-idem-')
        construir({p: os.path.join(clon, mapas.PLANES[p]['slug']) for p in mapas.PLANES}, datos, mes, prestamos,
                  log=lambda *_: None)
        difs = []
        for c in mapas.LIBROS:
            n = mapas.LIBROS[c][2]
            difs += diferencias(digest(os.path.join(carpetas[mapas.PLAN_DE[c]], n)),
                                digest(os.path.join(clon, mapas.PLANES[mapas.PLAN_DE[c]]['slug'], n)), n)
        shutil.rmtree(clon)
        idem = {'ejecutada': True, 'diferencias': len(difs), 'detalle': difs[:30]}
        print('  diferencias 1.ª vs 2.ª construcción: %d' % len(difs), flush=True)
        for d in difs[:10]:
            print('    ' + d)

    print('== 3/4 · inject_cache (AL FINAL, sobre los 4, en este mismo proceso)', flush=True)
    cache = inject_cache(rutas)

    print('== 4/4 · calibración D15 y cifras_caso.json', flush=True)
    fallos = []
    if idem.get('diferencias'):
        fallos.append('idempotencia: %d diferencias' % idem['diferencias'])
    fallos += ['inject_cache %s' % c['salida'] for c in cache if not c['ok'] or c['fallos_pycel'] != 0]
    calib = OrderedDict()
    for plan, corto in PLANES_CORTO.items():
        leer = Cache(os.path.join(carpetas[plan], mapas.LIBROS[corto][2]), corto)
        m, f = d15(leer, plan, datos)
        if leer('0. Supuestos', 'B31') != prestamos[plan]:
            f.append('B31 = %r ≠ préstamo calibrado %r' % (leer('0. Supuestos', 'B31'), prestamos[plan]))
        if prestamo_propuesto(leer, censo_libro(datos, corto)) != prestamos[plan]:
            f.append('el préstamo propuesto por Financing ya no es %r' % prestamos[plan])
        calib[plan] = OrderedDict([('metricas', m), ('fallos', f)])
        print('  %s · préstamo %s · ventas %.0f · DSCR mín %.2f · saldo mín %.0f · cobertura %.1f%% · holgura %.1f%% '
              '· margen bruto %.1f%% · neto a1 %.1f%% · payback %s%s' % (
                  plan, prestamos[plan], m['ventas_a1'], m['dscr_min'], m['saldo_minimo'], m['cobertura_horas'] * 100,
                  m['holgura_caja_pct'] * 100, m['margen_bruto_pct'] * 100, m['margen_neto_a1'] * 100, m['payback'],
                  '' if not f else ' · FALLOS: ' + '; '.join(f)))
        for rot, r in m['ratios_pyg'].items():
            print('      %-30s %s umbral %s %.3f → %s' % (rot, ['%.3f' % v for v in r['valores']], r['sentido'],
                                                          r['umbral'], ['OK' if x else 'ámbar' for x in r['cumple']]))
        fallos += ['D15 %s: %s' % (plan, x) for x in f]
    destino_json = os.path.join(AQUI, 'cifras_caso.json') if not args.dry_run else \
        os.path.join(os.path.dirname(carpetas['ft']), 'cifras_caso.json')
    escribir_cifras(carpetas, datos, destino_json)
    print('  cifras_caso.json → %s' % destino_json)

    salida = OrderedDict([('fecha', datetime.datetime.now().isoformat(timespec='seconds')),
                          ('modo', 'dry-run' if args.dry_run else 'real'), ('carpetas', carpetas), ('mes', mes),
                          ('prestamos', prestamos), ('libros', informe), ('idempotencia', idem),
                          ('inject_cache', cache), ('calibracion', calib), ('cifras', destino_json),
                          ('sha256', OrderedDict((os.path.basename(r), sha256(r)) for r in rutas)),
                          ('fallos', fallos), ('segundos', round(time.time() - t0, 1))])
    if args.json:
        with open(args.json, 'w', encoding='utf-8') as fh:
            json.dump(salida, fh, ensure_ascii=False, indent=1, default=str)
        print('informe → %s' % args.json)
    if fallos:
        print('\nFALLOS:\n  ' + '\n  '.join(fallos))
        sys.exit(1)
    print('\nOK: 4 xlsx en %s (%.0f s)' % (sorted(set(carpetas.values())), time.time() - t0))


if __name__ == '__main__':
    try:
        main()
    except Aborta as e:
        raise SystemExit('ABORTADO: %s' % e)
