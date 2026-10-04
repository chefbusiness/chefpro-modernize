#!/usr/bin/env python3
"""
ensamblar_docx.py — Food Truck + Coffee Shop Business Plan Kit (EN) · docx EN párrafo a párrafo (SPEC D22-D23, G8).

Copia el docx ES publicado (SOLO LECTURA) y, sin tocar el XML de formato:
  1. sustituye el texto del ÚNICO run de cada párrafo (conserva rPr y pPr): fijos de mapas (portada, aviso, índice,
     encabezados D22, cierre D23) + textos de los traductores (`docx_en_<plan>_a.json` + `_b.json`) con los tokens
     {{nombre}} rellenados desde `cifras_caso.json` (mapas.formatear);
  2. CAF: mueve el bloque financiero de §7 a §8 (mapas.DOCX_MOVER) y da al párrafo 134 el formato de un párrafo de
     cuerpo normal (mapas.DOCX_RPR_DE: en el ES era un aviso en rojo);
  3. core props EN (D20/D23: título, asunto, autor John Guerrero, palabras clave, URL, categoría, idioma en-US);
  4. página US Letter;
  5. guarda y valida G8 sobre el fichero guardado: mismo nº de párrafos y misma secuencia de estilos que el ES (tras
     la permutación declarada), mismo formato de run y de párrafo, encabezados D22 e índice exactos, cero español,
     «EUR», «€», «ChefBusiness», «chefbusiness.co», no latinos, tokens sin rellenar o «[Contenido pendiente»;
     importes «$» y % = cifras del caso (tokens) o datos con fuente de research/SPEC; core props; Letter.

    python3 ensamblar_docx.py --plan ft [--en A.json B.json] [--cifras cifras_caso.json] [--salida F.docx]
    python3 ensamblar_docx.py --plan caf --cifras-desde-es --en …        # cifras de PRUEBA desde la caché del xlsx ES
    python3 ensamblar_docx.py --autotest [--dir /tmp/bp-docx]             # pruebas sintéticas (ver abajo)

Por defecto: --en docx_en_<plan>_a.json docx_en_<plan>_b.json, --cifras cifras_caso.json y --salida
astro-site/public/dl/<slug>/<fichero EN> (mapas.DOCX). Sale ≠ 0 si G8 falla.

Formato de `cifras_caso.json` (lo escribe aplicar_en.py con mapas.calcular_cifras() sobre la caché del xlsx EN):
  {"meta": {"generado", "script", "fuentes": {"ft": "<xlsx EN>", "caf": "<xlsx EN>"}},
   "ft":  {"<token>": {"valor": 254800.0, "fmt": "usd", "texto": "$254,800", "hoja_es": "PyG 3 Años", "celda": "B9"},
           "<derivado>": {"valor": …, "fmt": …, "texto": …, "deriv": true}, …},
   "caf": {…}}
  El ensamblador usa «texto» si viene y, si no, formatea «valor» con el fmt de mapas.TOKENS_DOCX (manda mapas).

--autotest (en un directorio temporal; nunca escribe en astro-site/):
  A. sintético EN de los 2 planes (cada párrafo traducible = una frase inglesa con 2-3 tokens; cifras de prueba
     sacadas de la caché del xlsx ES) → G8 VERDE, formato idéntico, tokens rellenados, CAF movido;
  B. «ES con prefijo» (el texto ES con «[EN] » delante) → G8 debe dar ROJO por español;
  C. defectos inyectados en el EN sintético: token desconocido, importe sin fuente, «chefbusiness.co», párrafo que
     falta → cada uno debe abortar o dar ROJO.
"""
import argparse
import copy
import datetime
import json
import os
import shutil
import sys
import tempfile
from collections import OrderedDict

import docx
from docx.shared import Emu

AQUI = os.path.dirname(os.path.abspath(__file__))
REPO = os.path.abspath(os.path.join(AQUI, '..', '..', '..'))
sys.path.insert(0, AQUI)
import mapas                                                     # noqa: E402
import check_textos                                              # noqa: E402

LETTER = (7772400, 10058400)


class Aborta(Exception):
    pass


def cargar_es(plan):
    de = json.load(open(os.path.join(mapas.DATOS, 'docx_es_%s.json' % plan), encoding='utf-8'))
    ruta = os.path.join(REPO, mapas.PLANES[plan]['dir_es'], mapas.DOCX[plan][0])
    return de, ruta


def cargar_en(plan, rutas):
    en = OrderedDict()
    for r in rutas:
        d = json.load(open(r, encoding='utf-8'))
        for k, v in d.items():
            if k in en:
                raise Aborta('párrafo %s repetido en dos ficheros EN' % k)
            en[k] = v
    return en


def cifras_desde_es(plan):
    """Cifras de PRUEBA: caché del xlsx ES publicado (valores del caso español; solo para el autotest)."""
    import openpyxl
    censo = json.load(open(os.path.join(mapas.DATOS, 'censo_es.json'), encoding='utf-8'))
    corto = mapas.corto_plan(plan)
    wb = openpyxl.load_workbook(mapas.ruta_es(corto, REPO), data_only=True)
    return mapas.calcular_cifras(plan, lambda h, c: wb[h][c].value, censo['tokens_docx'][plan])


def textos_finales(plan, de, en, cifras):
    """Índice ES → texto EN final (fijos + traductores con tokens rellenados). Aborta ante cualquier hueco."""
    fijos = mapas.textos_fijos_docx(plan)
    trad = [k for k, e in de['parrafos'].items() if e.get('traducir')]
    falta = [k for k in trad if k not in en]
    sobra = [k for k in en if k not in trad]
    if falta:
        raise Aborta('faltan %d párrafos EN: %s' % (len(falta), ', '.join(falta[:20])))
    if sobra:
        raise Aborta('párrafos EN que no son traducibles (fijos o vacíos): %s' % ', '.join(sobra[:20]))
    validos = mapas.tokens_de(plan)

    def rellena(k, t):
        def sub(m):
            tok = m.group(1)
            if tok not in validos:
                raise Aborta('%s[%s]: token desconocido {{%s}}' % (plan, k, tok))
            c = cifras.get(tok)
            if c is None:
                raise Aborta('%s[%s]: {{%s}} no está en cifras_caso.json' % (plan, k, tok))
            fmt = validos[tok][0]
            return c.get('texto') or mapas.formatear(c['valor'], fmt)
        return mapas.RX_TOKEN.sub(sub, t)

    out = OrderedDict()
    for k, v in fijos.items():
        out[int(k)] = v
    for k in trad:
        out[int(k)] = rellena(k, en[k])
    return out


def poner_texto(p, texto, plan, i):
    runs = p.runs
    if len(runs) == 1:
        runs[0].text = texto
        return
    if not runs:
        raise Aborta('%s[%d]: párrafo sin runs con texto EN' % (plan, i))
    rprs = {(r._r.rPr.xml if r._r.rPr is not None else '') for r in runs}
    if len(rprs) > 1:
        raise Aborta('%s[%d]: %d runs con formatos distintos: no se puede sustituir sin perder formato (F2-NOTAS §4)'
                     % (plan, i, len(runs)))
    runs[0].text = texto                                    # todos los runs comparten formato: el primero lo lleva todo
    for r in runs[1:]:
        r._r.getparent().remove(r._r)


def ensamblar(plan, en, cifras, salida):
    de, ruta = cargar_es(plan)
    d = docx.Document(ruta)
    pars = list(d.paragraphs)                                # referencias ANTES de mover: índices = ES
    if len(pars) != de['meta']['n_parrafos']:
        raise Aborta('el docx ES ya no es el del censo (%d párrafos)' % len(pars))
    finales = textos_finales(plan, de, en, cifras)
    for i, t in finales.items():
        poner_texto(pars[i], t, plan, i)
    for i, j in mapas.DOCX_RPR_DE.get(plan, {}).items():
        r_dst, r_src = pars[i].runs[0]._r, pars[j].runs[0]._r
        if r_dst.rPr is not None:
            r_dst.remove(r_dst.rPr)
        if r_src.rPr is not None:
            r_dst.insert(0, copy.deepcopy(r_src.rPr))
    mov = mapas.DOCX_MOVER.get(plan)
    if mov:
        ancla = pars[mov['despues_de']]._p
        for i in mov['orden']:
            el = pars[i]._p
            ancla.addnext(el)
            ancla = el
    cp = d.core_properties
    dp = mapas.docprops_en(plan)
    cp.title, cp.subject, cp.keywords = dp['title'], dp['subject'], dp['keywords']
    cp.author, cp.last_modified_by, cp.category = 'John Guerrero', 'AI Chef Pro', dp['category']
    cp.comments = dp['description']
    cp.language = 'en-US'
    cp.revision = 1
    ahora = datetime.datetime.now(datetime.timezone.utc).replace(microsecond=0, tzinfo=None)
    cp.created = cp.modified = ahora
    for s in d.sections:
        s.page_width, s.page_height = Emu(LETTER[0]), Emu(LETTER[1])
    os.makedirs(os.path.dirname(os.path.abspath(salida)), exist_ok=True)
    d.save(salida)
    return salida


def orden_final(plan, n):
    """Permutación declarada: lista de índices ES en el orden en que quedan en el EN."""
    orden = list(range(n))
    mov = mapas.DOCX_MOVER.get(plan)
    if mov:
        resto = [i for i in orden if i not in mov['orden']]
        k = resto.index(mov['despues_de']) + 1
        orden = resto[:k] + mov['orden'] + resto[k:]
    return orden


def g8(plan, salida, cifras=None):
    """Valida el docx EN guardado contra el ES. Devuelve la lista de fallos (vacía = VERDE)."""
    fallos = []
    de, ruta = cargar_es(plan)
    es = docx.Document(ruta)
    ennd = docx.Document(salida)
    pe, pn = es.paragraphs, ennd.paragraphs
    if len(pe) != len(pn):
        return ['nº de párrafos %d ≠ ES %d' % (len(pn), len(pe))]
    orden = orden_final(plan, len(pe))
    rpr_de = mapas.DOCX_RPR_DE.get(plan, {})
    for pos, i in enumerate(orden):
        a, b = pe[i], pn[pos]
        if a.style.name != b.style.name:
            fallos.append('párrafo ES %d → EN %d: estilo %s ≠ %s' % (i, pos, b.style.name, a.style.name))
        ppr_a = a._p.pPr.xml if a._p.pPr is not None else ''
        ppr_b = b._p.pPr.xml if b._p.pPr is not None else ''
        if ppr_a != ppr_b:
            fallos.append('párrafo ES %d → EN %d: formato de párrafo (pPr) distinto' % (i, pos))
        if len(a.runs) != len(b.runs):
            fallos.append('párrafo ES %d → EN %d: %d runs ≠ %d' % (i, pos, len(b.runs), len(a.runs)))
            continue
        src = pe[rpr_de[i]] if i in rpr_de else a
        for ra, rb in zip(src.runs[:1] if i in rpr_de else a.runs, b.runs):
            xa = ra._r.rPr.xml if ra._r.rPr is not None else ''
            xb = rb._r.rPr.xml if rb._r.rPr is not None else ''
            if xa != xb:
                fallos.append('párrafo ES %d → EN %d: formato de run (rPr) distinto' % (i, pos))
        if a.text.strip() and not b.text.strip():
            fallos.append('párrafo ES %d → EN %d: se quedó vacío' % (i, pos))
        if not a.text.strip() and b.text.strip():
            fallos.append('párrafo ES %d → EN %d: un vacío del ES tiene texto' % (i, pos))
    # encabezados e índice (posición EN = posición ES: no se mueven)
    for n, (ih, ii) in enumerate(zip(mapas.DOCX_ENCABEZADOS[plan], mapas.DOCX_INDICE[plan]), 1):
        if pn[orden.index(ih)].text != '%d. %s' % (n, mapas.ENCABEZADOS_EN[n - 1]):
            fallos.append('encabezado %d: %r' % (n, pn[orden.index(ih)].text))
        if pn[orden.index(ii)].text != '%d.  %s' % (n, mapas.ENCABEZADOS_EN[n - 1]):
            fallos.append('índice %d: %r' % (n, pn[orden.index(ii)].text))
    # restos
    dol_ok, pct_ok = check_textos.cifras_permitidas()
    tok_dol, tok_pct = set(), set()
    for c in (cifras or {}).values():
        t = c.get('texto') or ''
        dl, pc = check_textos.cifras_texto(t)
        tok_dol |= {v for v, _ in dl}
        tok_pct |= {v for v, _ in pc}
    for pos, p in enumerate(pn):
        t = p.text
        if not t.strip():
            continue
        donde = 'EN[%d] (ES %d)' % (pos, orden[pos])
        for motivo, ctx in mapas.restos_espanol(t):
            fallos.append('%s: %s · %s' % (donde, motivo, ctx.strip()[:60]))
        if mapas.no_latinos(t):
            fallos.append('%s: no latinos %s' % (donde, mapas.no_latinos(t)))
        if '{{' in t or '}}' in t:
            fallos.append('%s: token sin rellenar' % donde)
        if 'pendiente' in t.lower() or '[Content' in t:
            fallos.append('%s: marcador de contenido pendiente' % donde)
        if check_textos.PROHIBIDO.search(t):
            fallos.append('%s: afirmación prohibida' % donde)
        dl, pc = check_textos.cifras_texto(t)
        for v, txt in dl:
            if v not in dol_ok and v not in tok_dol:
                fallos.append('%s: importe «%s» que no es del caso ni de research/SPEC' % (donde, txt))
        for v, txt in pc:
            if v not in pct_ok and v not in tok_pct:
                fallos.append('%s: porcentaje «%s» que no es del caso ni de research/SPEC' % (donde, txt))
    # core props, idioma y página
    cp = ennd.core_properties
    dp = mapas.docprops_en(plan)
    for campo, esperado in (('title', dp['title']), ('subject', dp['subject']), ('keywords', dp['keywords']),
                            ('author', 'John Guerrero'), ('comments', dp['description']),
                            ('category', dp['category']), ('language', 'en-US')):
        if getattr(cp, campo) != esperado:
            fallos.append('core props %s = %r' % (campo, getattr(cp, campo)))
    for campo in ('title', 'subject', 'keywords', 'comments', 'category', 'author', 'last_modified_by'):
        v = getattr(cp, campo) or ''
        if mapas.restos_espanol(v) or mapas.RX_MARCA_ES.search(v) or 'python-docx' in v:
            fallos.append('core props %s con español, marca ES o «python-docx»: %r' % (campo, v))
    for s in ennd.sections:
        if (s.page_width, s.page_height) != LETTER:
            fallos.append('página no Letter: %s × %s' % (s.page_width, s.page_height))
    return fallos


# --------------------------------------------------------------------------
# autotest
# --------------------------------------------------------------------------
FRASE = ('Synthetic paragraph {i} of the {plan} plan for the assembly test. Year 1 revenue is {{{{ventas_a1}}}} and '
         'the plan breaks even at {{{{equilibrio_clientes_dia}}}} customers a day, with a lowest DSCR of '
         '{{{{dscr_min}}}}.')
SUBT = 'Synthetic subtitle {i}'


def en_sintetico(plan, de):
    out = OrderedDict()
    for k, e in de['parrafos'].items():
        if e.get('traducir'):
            out[k] = (SUBT if e['rol'] == 'subtitulo' else FRASE).format(i=k, plan=plan)
    return out


def autotest(dirbase):
    os.makedirs(dirbase, exist_ok=True)
    res = []

    def caso(nombre, ok, detalle=''):
        res.append((nombre, ok, detalle))
        print('%s %s %s' % ('OK   ' if ok else 'FALLO', nombre, detalle))

    for plan in mapas.PLANES:
        de, _ = cargar_es(plan)
        cif = cifras_desde_es(plan)
        en = en_sintetico(plan, de)
        # A. sintético → VERDE
        sal = os.path.join(dirbase, 'A-%s.docx' % plan)
        ensamblar(plan, en, cif, sal)
        f = g8(plan, sal, cif)
        caso('A %s sintético: G8 VERDE' % plan, not f, '; '.join(f[:5]))
        dn = docx.Document(sal)
        orden = orden_final(plan, len(dn.paragraphs))
        k0 = next(k for k, e in de['parrafos'].items() if e.get('traducir') and e['rol'] == 'cuerpo')
        txt = dn.paragraphs[orden.index(int(k0))].text
        caso('A %s tokens rellenados' % plan, '{{' not in txt and cif['ventas_a1']['texto'] in txt, txt[:90])
        if plan == 'caf':
            pos110 = orden.index(110)
            seq = [dn.paragraphs[pos110 + j].text[:30] for j in range(3)]
            caso('A caf bloque financiero en §8 (tras «8. Financial Plan» y su separador)',
                 dn.paragraphs[pos110 - 2].text == '8. Financial Plan' and orden[pos110 + 1] == 134, str(seq))
            r134 = dn.paragraphs[orden.index(134)].runs[0]
            col = r134.font.color.rgb if r134.font.color is not None and r134.font.color.type is not None else None
            caso('A caf 134 sin rojo (formato del 112)', str(col) != 'CC0000', 'color=%s' % col)
        # B. ES con prefijo → ROJO por español
        en_es = OrderedDict((k, '[EN] ' + de['parrafos'][k]['texto']) for k in en)
        sal_b = os.path.join(dirbase, 'B-%s.docx' % plan)
        ensamblar(plan, en_es, cif, sal_b)
        fb = g8(plan, sal_b, cif)
        caso('B %s ES con prefijo: G8 ROJO por español' % plan, any('ES' in x or 'tilde' in x for x in fb),
             '%d fallos, p. ej. %s' % (len(fb), fb[:1]))
        # C. defectos inyectados
        defectos = [
            ('token desconocido', lambda d: d.update({k0: d[k0] + ' {{no_existe}}'}), 'aborta'),
            ('importe sin fuente', lambda d: d.update({k0: d[k0] + ' It costs $123,457.'}), 'rojo'),
            ('marca ES', lambda d: d.update({k0: d[k0] + ' See chefbusiness.co.'}), 'rojo'),
            ('párrafo que falta', lambda d: d.pop(k0), 'aborta'),
            ('pendiente', lambda d: d.update({k0: '[Contenido pendiente de generacion]'}), 'rojo'),
        ]
        for nombre, mut, esperado in defectos:
            d2 = OrderedDict(en)
            mut(d2)
            sal_c = os.path.join(dirbase, 'C-%s.docx' % plan)
            try:
                ensamblar(plan, d2, cif, sal_c)
                fc = g8(plan, sal_c, cif)
                caso('C %s %s → %s' % (plan, nombre, esperado), esperado == 'rojo' and bool(fc),
                     '%d fallos' % len(fc))
            except Aborta as ex:
                caso('C %s %s → %s' % (plan, nombre, esperado), esperado == 'aborta', str(ex)[:80])
    n_ok = sum(1 for r in res if r[1])
    print('autotest: %d/%d' % (n_ok, len(res)))
    return 0 if n_ok == len(res) else 1


def main():
    ap = argparse.ArgumentParser(description=__doc__.split('\n')[1])
    ap.add_argument('--plan', choices=list(mapas.PLANES))
    ap.add_argument('--en', nargs='+')
    ap.add_argument('--cifras', default=os.path.join(mapas.DATOS, 'cifras_caso.json'))
    ap.add_argument('--cifras-desde-es', action='store_true', help='cifras de PRUEBA desde la caché del xlsx ES')
    ap.add_argument('--salida')
    ap.add_argument('--autotest', action='store_true')
    ap.add_argument('--dir', default=os.path.join(tempfile.gettempdir(), 'bp-docx-autotest'))
    args = ap.parse_args()
    if args.autotest:
        return autotest(args.dir)
    if not args.plan:
        ap.error('--plan %s (o --autotest)' % '|'.join(mapas.PLANES))
    plan = args.plan
    rutas = args.en or [os.path.join(mapas.DATOS, 'docx_en_%s_%s.json' % (plan, s)) for s in ('a', 'b')]
    try:
        en = cargar_en(plan, rutas)
        if args.cifras_desde_es:
            cif = cifras_desde_es(plan)
        else:
            cif = json.load(open(args.cifras, encoding='utf-8'))[plan]
        salida = args.salida or os.path.join(REPO, mapas.PLANES[plan]['dir_en'], mapas.DOCX[plan][1])
        ensamblar(plan, en, cif, salida)
    except Aborta as ex:
        print('ABORTA', ex)
        return 2
    f = g8(plan, salida, cif)
    for x in f:
        print('G8 FALLO', x)
    print('ensamblar_docx %s → %s · G8 %s' % (plan, os.path.relpath(salida, REPO) if salida.startswith(REPO) else salida,
                                              'VERDE' if not f else 'ROJO (%d)' % len(f)))
    return 1 if f else 0


if __name__ == '__main__':
    sys.exit(main())
