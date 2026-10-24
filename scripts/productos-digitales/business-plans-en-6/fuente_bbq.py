#!/usr/bin/env python3
"""
fuente_bbq.py — material de CONTENIDO para los redactores: el checklist ES del Parrillero / Asador 2.2 (83 tareas, desde
`../parrillero-2_2/datos_checklist.py`, sin abrir el xlsx), los párrafos ES 2.2 del plan en las 61 ranuras del docx EN
(`../parrillero-2_2/textos/par_es_{a,b}.json`, con sus {TOKENS}), y las tablas y textos del manual técnico, la carta de
cortes, las 10 experiencias y el contrato (python-docx) en `fuente_es_bbq.json`. El circuito de texto corre sobre el FOOD
TRUCK ES (las ranuras); el contenido de BBQ sale de aquí y de research (D101, D107, D108). Sus CIFRAS no se usan (D110).
Corre en el VPS (python-docx). Más simple que `../business-plans-en-5/fuente_mbar.py`. Sesión Claude Code, 24-oct-2026
(F2, paso 1).

    python3 bp6.py fuente
"""
import datetime
import hashlib
import importlib.util
import json
import os
from collections import OrderedDict

import docx

AQUI = os.path.dirname(os.path.abspath(__file__))
REPO = os.path.abspath(os.path.join(AQUI, '..', '..', '..'))
PAR = os.path.join(os.path.dirname(AQUI), 'parrillero-2_2')
DIR = os.path.join(REPO, 'astro-site', 'public', 'dl', 'plan-negocio-parrillero-asador-eventos')
MANUAL = 'manual-tecnico-parrillero-profesional.docx'
CARTA = 'carta-12-cortes-carne.docx'
EXPERIENCIAS = '10-experiencias-tematicas-eventos.docx'
CONTRATO = 'modelos-contrato-parrillero-asador-eventos.docx'
# F1-inventario §5: tareas solo España (fase, nº de tarea 1-based) — fuera del EN; y las parciales (lo español, fuera)
SOLO_ES = {(1, 3), (1, 4), (1, 5), (1, 8), (1, 10), (1, 11), (2, 3), (2, 5), (4, 2), (4, 7), (5, 14)}
PARCIALES = {(1, 7): 'Veri*factu', (2, 7): 'MER / vértebra', (2, 12): 'ADR', (5, 10): 'Hosply', (6, 13): 'FACe'}
FASE_EN = OrderedDict([
    (1, 'P1 Business Setup (1-2 → LLC; 6 banco; 7 y 9 → contador y estimated taxes; 12 → contrato y cancelación)'),
    (2, 'P2 Pits, Fire & Permits (1 commissary; 2 → county catering / health permit; 8-9 → GL, commercial auto + '
        'remolque; 11-15 → fire marshal, propane backup, extintores, burn bans, cenizas; 16 → temporary event permits) · '
        'P1 (4 → plan HACCP; 10 → fuerza mayor) · P4 (6 alérgenos)'),
    (3, 'P3 Equipment (1-5 → smoker trailer, live-fire rig, propane backup, leña; 6, 8-9 → smallwares, hot boxes, '
        'probes, básculas; 12 → meat purveyor USDA-inspected + Restaurant Depot; 13 sides) · P2 (14 → van) · P4 (11)'),
    (4, 'P4 Staff (1 → W-2; 3-5 → food handler cards, alérgenos, OSHA calor / humo / quemaduras; 6 → registros FLSA; '
        '8 → staffing por invitados; 9 → guion del live-fire / carving; 10 → fichas por paquete)'),
    (5, 'P5 Marketing (1 → USPTO; 3-7, 11-13, 15-16; 8 WhatsApp → texto / SMS; 9 → The Knot / WeddingWire; '
        '10 → ezCater)'),
    (6, 'P6 First 90 Days (1-3 briefing, visita y depósito; 5-7 alérgenos, plan B de lluvia y burn ban, pedido de carne '
        '5-7 días; 8 rubs y porcionado en la commissary; 9-15 → carga, encendido, cocción con sonda, cenizas, factura, '
        'reseña, yield real frente a la calculadora)'),
])
USO = ('CONTENIDO de BBQ catering para los redactores (D101, D107, D108). NO se traduce tal cual: se condensa y se adapta '
       'a EE. UU. con research (F1-research-us-delta.md §1, §5, §6, §8 y el glosario §9). Las CIFRAS del ES NO se usan '
       '(D110: 190.000, 95 eventos/año, 80 invitados, 25 €/invitado, 31,8 %…): las del caso salen del libro EN (tokens). '
       'Fuera: tareas marcadas SOLO_ES y lo español de las PARCIALES, «España», «ChefBusiness», normativa española (RD, '
       'Reglamento (CE), AESAN, APPCC, CCAA), proveedores y directorios españoles, importes en €, «autónomo», «fijo '
       'discontinuo» y «freelance» (en US: part-time hourly W-2 staff), el anexo CCAA y las erratas del ES (D113: merma '
       'del brisket 35-40 % frente a 40-50 %; «18 cortes» son 16 de carne; la «carta 12 cortes» son 10; [98] con los días '
       'de cobro en los pagos). Temperaturas en °F: la de calidad y el mínimo del FDA Food Code, siempre distinguidas')


def sha(p):
    return hashlib.sha256(open(p, 'rb').read()).hexdigest()


def checklist():
    p = os.path.join(PAR, 'datos_checklist.py')
    spec = importlib.util.spec_from_file_location('datos_checklist_par', p)
    m = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(m)
    fases = OrderedDict()
    for k, (hoja, titulo, tareas) in enumerate(m.FASES, 1):
        lista = []
        for j, (t, resp, plazo, nota) in enumerate(tareas, 1):
            x = OrderedDict([('n', j), ('tarea', t), ('responsable', resp), ('plazo', plazo), ('nota', nota)])
            if (k, j) in SOLO_ES:
                x['marca'] = 'SOLO_ES (fuera)'
            elif (k, j) in PARCIALES:
                x['marca'] = 'PARCIAL: fuera lo español (%s)' % PARCIALES[(k, j)]
            lista.append(x)
        fases[hoja] = OrderedDict([('titulo', titulo), ('fase_en_sugerida', FASE_EN[k]), ('tareas', lista)])
    return p, m.TITULO, fases


def ranuras():
    out, fuentes = OrderedDict(), []
    for f in ('a', 'b'):
        p = os.path.join(PAR, 'textos', 'par_es_%s.json' % f)
        fuentes.append(p)
        out.update(json.load(open(p, encoding='utf-8')))
    return fuentes, OrderedDict(sorted(out.items(), key=lambda kv: int(kv[0])))


def parrafos(nombre, minimo=3):
    p = os.path.join(DIR, nombre)
    d = docx.Document(p)
    pars = [OrderedDict([('i', i), ('estilo', x.style.name), ('texto', x.text)]) for i, x in enumerate(d.paragraphs)
            if len(x.text.strip()) >= minimo and not set(x.text.strip()) <= set('━─—-· ')]
    tablas = [[[c.text for c in row.cells] for row in t.rows] for t in d.tables]
    return p, OrderedDict([('parrafos', pars), ('tablas', tablas)])


def main():
    pc, titulo, fases = checklist()
    pr, docx_es = ranuras()
    docs = OrderedDict()
    fuentes = [pc] + pr
    for clave, nombre in (('manual_tecnico', MANUAL), ('carta_cortes', CARTA), ('experiencias_10', EXPERIENCIAS),
                          ('contrato', CONTRATO)):
        p, contenido = parrafos(nombre)
        docs[clave] = contenido
        fuentes.append(p)
    out = OrderedDict([
        ('meta', OrderedDict([
            ('generado', datetime.datetime.now(datetime.timezone.utc).isoformat(timespec='seconds')),
            ('script', 'scripts/productos-digitales/business-plans-en-6/fuente_bbq.py'),
            ('uso', USO),
            ('fuentes', OrderedDict((os.path.relpath(x, REPO), sha(x)) for x in fuentes)),
        ])),
        ('checklist_es', OrderedDict([('titulo', titulo), ('fases', fases)])),
        ('docx_es_por_ranura', docx_es),
    ] + list(docs.items()))
    dest = os.path.join(AQUI, 'fuente_es_bbq.json')
    with open(dest, 'w', encoding='utf-8') as fh:
        json.dump(out, fh, ensure_ascii=False, indent=1)
        fh.write('\n')
    n = sum(len(f['tareas']) for f in fases.values())
    fuera = sum(1 for f in fases.values() for t in f['tareas'] if t.get('marca', '').startswith('SOLO_ES'))
    print('fuente_es_bbq.json: %d fases, %d tareas (%d fuera, %d parciales) · docx %d ranuras · %s' % (
        len(fases), n, fuera, len(PARCIALES), len(docx_es),
        ' · '.join('%s %d párrafos / %d tablas' % (k, len(v['parrafos']), len(v['tablas'])) for k, v in docs.items())))
    return 0


if __name__ == '__main__':
    raise SystemExit(main())
