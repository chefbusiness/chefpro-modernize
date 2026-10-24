#!/usr/bin/env python3
"""
calculadora_bbq.py — BBQ Catering Pricing Calculator (Pounds of Meat per Guest to Quote + Smoke Schedule & Shopping List),
D105-D106 del BBQ Catering Business Plan Kit (EN). Esqueleto y autotest de `../business-plans-en-5/calculadora_mbar.py`;
coste → precio de `../business-plans-en-4/calculadora_cat.py` (billable guests, service charge, interruptor del sales tax
sobre el service charge, depósito); escalado por libras y tramos de cancelación sin huecos de
`../parrillero-2_2/calculadora_par_es.py` (D105). Mismas reglas: entradas verdes con validación, salidas calculadas, el
margen se CALCULA, ejemplos con las mismas fórmulas, docProps D20, US Letter, sin anclas de precio, todo valor por defecto
«Kit estimate» salvo cita. Cinco pestañas:

  1. Pricing Calculator            paquete (drop-off · full-service BBQ buffet · live-fire premium) + interruptor «live pit
                                   on site» ($650 por defecto: eventective.com, Owen's BBQ); invitados y guest minimum;
                                   oz cocinadas por invitado por paquete (6 / 8 / 8); 5 líneas de carne (corte de la tabla
                                   de rendimientos, % de la mezcla, $/lb crudo EN BLANCO: lo teclea el usuario); sides,
                                   pan y salsas, desechables, leña; personal; alquileres, millas; margen objetivo, service
                                   charge, sales tax + interruptor, depósito del 50 % y 4 tramos de cancelación → libras
                                   crudas y piezas por línea, coste, precio por invitado y por lb cocinada, total, depósito,
                                   saldo, retención por tramo y margen calculado. Mientras falten precios de carne, el coste
                                   de comida usa el Kit estimate por invitado del paquete (research §5) y lo dice.
  2. Meat Yield & Smoke Schedule   13 cortes: peso crudo por pieza, temperatura del pit °F (°C), temperatura de CALIDAD
                                   (Kit estimate) y MÍNIMO del FDA Food Code en columnas distintas, horas, reposo,
                                   rendimiento (brisket y pork butt ≈ 50 %: fieryfoodscentral.com; el resto, Kit estimate),
                                   raciones por lb cruda y «Start the pit by» por fórmula.
  3. Shopping List by Guest Count  los invitados de la hoja 1 + 40 / 80 / 150 / 300 (tabla T4 del ES): libras y piezas por
                                   línea de carne, sides, rolls, salsa, leña, butcher paper, hotel pans, hot boxes,
                                   pit crew y servers.
  4. 3 Example Quotes              drop-off corporativo 40 inv. · graduation party staffed buffet 70 · boda 130 full-service
                                   + live pit on site: a ±10 % del precio de su línea en la mezcla de research §8 ($22 /
                                   $35 / $52), sin ningún precio de carne tecleado.
  5. Instructions

Corre en el VPS (openpyxl + pycel). Sesión Claude Code, 24-oct-2026 (F2, paso 1).

    python3 bp6.py calculadora              # genera + inject_cache + autotest
    python3 bp6.py calculadora --autotest   # solo el autotest sobre el fichero ya generado
"""
import argparse
import math
import os
import re
import sys
from collections import OrderedDict

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill, Protection
from openpyxl.worksheet.datavalidation import DataValidation

AQUI = os.path.dirname(os.path.abspath(__file__))
PD = os.path.dirname(AQUI)
REPO = os.path.abspath(os.path.join(AQUI, '..', '..', '..'))
sys.path.insert(0, os.path.join(PD, 'business-plans-en'))
sys.path.insert(0, PD)
os.environ.setdefault('BP_CONJUNTO', 'bbq')
import mapas as M                                                                  # noqa: E402

SLUG = 'bbq-catering-business-plan'
FICHERO = 'bbq-catering-pricing-calculator.xlsx'
TITULO = 'BBQ Catering Pricing Calculator (Pounds of Meat per Guest to Quote + Smoke Schedule & Shopping List)'
DESTINO = os.path.join(REPO, 'astro-site', 'public', 'dl', SLUG, FICHERO)
H1, H2, H3, H4, H5 = ('Pricing Calculator', 'Meat Yield & Smoke Schedule', 'Shopping List by Guest Count',
                      '3 Example Quotes', 'Instructions')
HOJAS_CALCULO = (H1, H2, H3, H4)
VERDE = PatternFill('solid', fgColor='E8F5E9')
GRIS = PatternFill('solid', fgColor='F2F2F2')
F = 'Calibri'
MONEDA, ENTERO, PCT, PCT1, DEC1, DEC2, HORA = '#,##0.00', '#,##0', '0%', '0.0%', '#,##0.0', '#,##0.00', 'h:mm AM/PM'

DROP = 'Drop-off (delivered & set, disposable pans)'
BUFFET = 'Full-service BBQ buffet (staffed, carving station)'
LIVE = 'Live-fire premium (Argentine asado / Santa Maria / whole hog)'
# paquete → (oz cocinadas por invitado, invitados por server / attendant (0 = ninguno), invitados por persona del pit
# crew, coste de comida por invitado mientras falten precios de carne, nota). oz: decisión SPEC D105 (6 / 8 / 8, dentro
# de «one-third to one-half pound» de hutchinsbbq.com y de los 6-8 oz de sonnybryans.com); servers: Kit estimate (1 por
# 25-30 en buffet); coste por invitado: Kit estimates de research §5 ($6.50 drop-off · $8 buffet · $17 live-fire)
PAQUETES = OrderedDict([
    (DROP, (6.0, 0, 150, 6.50, 'Delivered and set in disposable pans, no servers on site. 6 oz cooked is inside the '
                                '6-8 oz of sonnybryans.com and the one-third to one-half pound of hutchinsbbq.com.')),
    (BUFFET, (8.0, 30, 100, 8.00, 'Staffed buffet with a carving station: about 1 server or buffet attendant per 25-30 '
                                  'guests (Kit estimate). 8 oz = half a pound of cooked meat per guest.')),
    (LIVE, (8.0, 25, 50, 17.00, 'Argentine asado, Santa Maria tri-tip or whole hog cooked at the event: more pit crew '
                                'per guest and pricier cuts. Pick the cuts in the meat lines.')),
])
NOMBRES_PAQ = list(PAQUETES)

# --------------------------------------------------------------------------
# 2. Meat Yield & Smoke Schedule: 13 cortes. (corte, peso crudo por pieza lb, temperatura del pit °F, temperatura de
#    CALIDAD °F [Kit estimate], mínimo del FDA Food Code 2022 (texto), horas de cocción, horas de reposo y hold,
#    rendimiento cocinado ÷ crudo, fuente del rendimiento, nombre de la pieza). Pesos, temperaturas y horas = Kit estimate
# --------------------------------------------------------------------------
_FC_ASADO = '145 °F (63 °C) for 4 min as a whole roast; 155 °F (68 °C) for 17 s if injected'
_FC_INTACTA = '145 °F (63 °C) for 15 s; 155 °F (68 °C) for 17 s if injected or tenderized'
_FC_PICADA = '155 °F (68 °C) for 17 s (ground meat, sausage included)'
_FC_AVES = '165 °F (74 °C), instantaneous (poultry)'
_KIT = 'Kit estimate'
_LEIDA = 'fieryfoodscentral.com (about 50% of the fresh weight)'
CORTES = [
    ('Packer brisket (whole, untrimmed)', 14.0, 250, 203, _FC_ASADO, 12.0, 2.0, 0.50, _LEIDA, 'packers'),
    ('Pork butt (bone-in shoulder)', 8.0, 250, 200, _FC_ASADO, 10.0, 1.5, 0.50, _LEIDA, 'butts'),
    ('St. Louis ribs', 3.0, 250, 198, _FC_INTACTA, 5.5, 0.5, 0.65, _KIT, 'racks'),
    ('Baby back ribs', 2.25, 250, 195, _FC_INTACTA, 4.5, 0.5, 0.60, _KIT, 'racks'),
    ('Beef plate ribs (3-bone)', 5.5, 275, 203, _FC_INTACTA, 9.0, 1.0, 0.55, _KIT, 'plates'),
    ('Chicken thighs (bone-in)', 0.33, 300, 185, _FC_AVES, 1.75, 0.25, 0.68, 'Kit estimate (65-70%)', 'thighs'),
    ('Chicken halves', 2.0, 300, 175, _FC_AVES, 2.5, 0.25, 0.65, 'Kit estimate (65-70%)', 'halves'),
    ('Turkey breast (bone-in)', 7.0, 275, 165, _FC_AVES, 4.0, 0.5, 0.70, _KIT, 'breasts'),
    ('Sausage (smoked links)', 0.25, 275, 165, _FC_PICADA, 1.5, 0.25, 0.88, 'Kit estimate (85-90%)', 'links'),
    ('Tri-tip (Santa Maria, over red oak)', 2.5, 375, 145, _FC_ASADO, 1.0, 0.25, 0.75, _KIT, 'tri-tips'),
    ('Picanha (top sirloin cap)', 3.0, 400, 145, _FC_ASADO, 1.0, 0.25, 0.75, _KIT, 'caps'),
    ('Tomahawk / bone-in ribeye', 3.0, 250, 145, _FC_INTACTA, 1.5, 0.25, 0.70, _KIT, 'steaks'),
    ('Whole hog (optional, dressed)', 120.0, 250, 200, _FC_ASADO, 14.0, 1.0, 0.35, _KIT, 'hogs'),
]
NOMBRES_CORTE = [c[0] for c in CORTES]
Y_PAR = OrderedDict([('service', ('Service time on event day', HORA, 0.75, 'Type a time such as 6:00 PM.')),
                     ('cushion', ('Safety cushion before service (hours)', DEC1, 1.0,
                                  'Extra time for a slow cook or the stall. Kit estimate.')),
                     ('serving_oz', ('Cooked oz per serving of one meat', DEC1, 4.0,
                                     'With 2 meats on the plate, 4 oz each = half a pound of cooked meat per guest.'))])
Y_F0 = 10                                                              # primera fila de cortes en la hoja 2
Y_COL = OrderedDict([('peso', 'B'), ('pit', 'C'), ('pit_c', 'D'), ('cal', 'E'), ('cal_c', 'F'), ('fc', 'G'), ('horas', 'H'),
                     ('reposo', 'I'), ('rend', 'J'), ('fuente', 'K'), ('raciones', 'L'), ('coc_pieza', 'M'),
                     ('inicio', 'N'), ('dia', 'O')])
Y_PAR_FILA = {k: 5 + i for i, k in enumerate(Y_PAR)}                   # parámetros de la hoja 2 en B5:B7
T_CORTES = "'%s'!$A$%d:$A$%d" % (H2, Y_F0, Y_F0 + len(CORTES) - 1)


def _ycol(col):
    return "'%s'!$%s$%d:$%s$%d" % (H2, col, Y_F0, col, Y_F0 + len(CORTES) - 1)


# --------------------------------------------------------------------------
# 1. Pricing Calculator: entradas (clave, rótulo, formato, valor por defecto, nota). Fila = 5 + índice
# --------------------------------------------------------------------------
MEZCLA_DEF = [('Packer brisket (whole, untrimmed)', 0.35), ('Pork butt (bone-in shoulder)', 0.30), ('St. Louis ribs', 0.15),
              ('Chicken thighs (bone-in)', 0.15), ('Sausage (smoked links)', 0.05)]
_CARNE = []
for _i in range(1, 6):
    _CARNE += [
        ('cut%d' % _i, 'Meat line %d: cut' % _i, '@', MEZCLA_DEF[_i - 1][0],
         'Choose from the cuts on the "%s" tab (its yield and piece weight are used here).' % H2 if _i == 1 else
         'Choose a cut from the list.'),
        ('share%d' % _i, 'Meat line %d: share of the cooked meat' % _i, PCT, MEZCLA_DEF[_i - 1][1],
         'The five shares must add up to 100% (check in the results). 0% = line not used.' if _i == 1 else ''),
        ('price%d' % _i, 'Meat line %d: your raw price per lb' % _i, MONEDA, None,
         'Left blank on purpose: type your supplier\'s price per pound of RAW meat. This kit does not guess meat prices.'
         if _i == 1 else 'Your supplier\'s price per raw lb.'),
    ]
ENTRADAS = [
    ('package', 'Package', '@', BUFFET,
     'Choose from the list (table at the bottom of this tab): it sets the cooked meat per guest, the suggested staff '
     'and the food cost used while your meat prices are blank.'),
    ('guests', 'Guests expected', ENTERO, 75, 'Whole number. You cook and buy for these guests.'),
    ('minimum', 'Guest minimum in your contract', ENTERO, 30,
     'If fewer guests confirm, you bill the minimum (billable guests in the results). Operators read for this kit use '
     '30 (eventective.com, Owen\'s BBQ) and 75 for full catering (hutchinsbbq.com). 0 = no minimum.'),
] + _CARNE + [
    ('sides', 'Sides per guest (mac and cheese, beans, slaw, potato salad)', MONEDA, 2.00,
     'Kit estimate. Used together with your meat prices; while they are blank, the package food cost replaces it.'),
    ('bread', 'Bread, pickles, onions and BBQ sauce per guest', MONEDA, 0.60, 'Kit estimate. Same rule as the sides.'),
    ('disposables', 'Disposables per guest (pans, plates, cutlery, butcher paper, foil)', MONEDA, 0.75, 'Kit estimate.'),
    ('wood_ratio', 'Wood or charcoal per raw lb of meat (lb)', DEC2, 0.6,
     'Kit estimate: a full cook burns about 60-100 lb of wood. Pellets or charcoal: use your own ratio.'),
    ('wood_price', 'Wood or charcoal cost per lb', MONEDA, 0.15, 'Kit estimate: use your supplier\'s price.'),
    ('live_pit', 'Live pit on site (smoker trailer at the event)?', '@', 'No',
     'Yes = you tow the smoker trailer and cook or carve at the event, for a flat fee (next row). Ask your county first: '
     'cooking at the event can make the trailer a mobile food unit.'),
    ('live_pit_fee', 'Live pit on site fee (flat, per event)', MONEDA, 650.0,
     'Charged on top of the per-guest price when the switch above is Yes. A Los Angeles operator charges $650 for the '
     'smoker on site (eventective.com, Owen\'s BBQ).'),
    ('pit_crew', 'Pit crew on event day (pitmasters, pit cooks)', ENTERO, 1,
     'Compare with the suggested number in the results.'),
    ('servers', 'Servers, buffet attendants and carvers', ENTERO, 3,
     'Compare with the suggested number in the results.'),
    ('event_hours', 'Hours each on event day (loading, travel, setup, service, breakdown)', DEC1, 6.0, 'Kit estimate.'),
    ('cook_hours', 'Overnight cook and prep hours (pit cook, total)', DEC1, 8.0,
     'Trimming, rubs and the overnight cook for this event. One pit cook can watch several cooks at once: count only the '
     'hours this event takes. Kit estimate.'),
    ('wage_pit', 'Pit crew hourly wage', MONEDA, 21.0,
     'Kit estimate, about the 75th percentile of restaurant cooks (BLS OEWS, May 2025: $21.45). W-2 employees.'),
    ('wage_server', 'Server / attendant hourly wage', MONEDA, 18.0,
     'Kit estimate, about the 75th percentile of nonrestaurant food servers (BLS OEWS, May 2025: $18.51). W-2.'),
    ('payroll', "Employer payroll taxes and workers' comp (% of wages)", PCT, 0.10, 'Kit estimate.'),
    ('rentals', 'Rentals passed through (tables, linens, chafing dishes, tent)', MONEDA, 0.0,
     'What the rental company charges you for this event.'),
    ('miles', 'Round-trip miles', ENTERO, 30, 'Commissary → venue → commissary.'),
    ('per_mile', 'Cost per mile (van, or pickup + smoker trailer)', MONEDA, 0.70,
     'Kit estimate: use your real cost per mile or the current IRS standard mileage rate.'),
    ('event_fee', 'Flat event or travel fee charged to the client', MONEDA, 0.0,
     'Optional flat charge on top of the per-guest price. 0 = none.'),
    ('commission', 'Marketplace fee on this booking (% of subtotal)', PCT, 0.0,
     'Corporate catering marketplaces such as ezCater keep about 15% (Kit estimate). 0 = direct booking.'),
    ('margin', 'Target gross margin on the food & service subtotal', PCT, 0.45,
     'Share of the subtotal left after the direct costs above and the marketplace fee. Kit estimate.'),
    ('sc', 'Service charge (% of subtotal)', PCT, 0.0,
     'Added on top. The part you pay out to staff is wages, not tips (IRS Rev. Rul. 2012-18), with payroll taxes; the '
     'part you keep is revenue.'),
    ('tax', 'Sales tax rate (state + local)', PCT1, 0.08, 'Kit estimate: use your own state + local rate.'),
    ('tax_sc', 'Does your state tax the service charge?', '@', 'Yes',
     'Sales tax always applies to the food & service subtotal here. CA: a mandatory service charge is taxable. NY: not '
     'taxable only if it is separately stated, identified as a gratuity and all of it goes to the staff. TX: not taxable '
     'up to 20% if separately stated and paid to the staff. Check your state.'),
    ('deposit', 'Nonrefundable deposit to book the date (% of total)', PCT, 0.50,
     'A 50% nonrefundable deposit is what the operators read for this kit ask (eventective.com, hutchinsbbq.com); the '
     'balance is due before the event.'),
    ('d1', 'Cancellation tier 1 starts at (days before the event)', ENTERO, 60,
     'Tiers run with no gaps: tier 1 = this many days or more; tier 2 = from the next threshold up to here; and so on.'),
    ('d2', 'Cancellation tier 2 starts at (days before the event)', ENTERO, 30, ''),
    ('d3', 'Cancellation tier 3 starts at (days before the event)', ENTERO, 14, 'Tier 4 = fewer days than this, or a no-show.'),
    ('r1', 'Share of the total you keep in tier 1', PCT, 0.50, 'Kit estimate: the nonrefundable deposit.'),
    ('r2', 'Share of the total you keep in tier 2', PCT, 0.75, 'Kit estimate.'),
    ('r3', 'Share of the total you keep in tier 3', PCT, 0.90, 'Kit estimate.'),
    ('r4', 'Share of the total you keep in tier 4', PCT, 1.00, 'Kit estimate: the meat is already bought.'),
]
K = {k: 5 + i for i, (k, *_r) in enumerate(ENTRADAS)}                  # clave → fila de la hoja 1
F_CALC = 5 + len(ENTRADAS) + 1                                         # cabecera del bloque de cálculo
_LINEAS = range(1, 6)
SALIDAS = [
    ('billable', 'Billable guests', ENTERO, 'max({guests},{minimum})', 'The larger of guests expected and the minimum.'),
    ('oz_guest', 'Cooked meat per guest for this package (oz)', DEC1, 'paq:2', 'From the package table below.'),
    ('cooked_lb', 'Cooked meat to serve (lb)', DEC1, '{guests}*{oz_guest}/16', 'Guests expected × oz per guest ÷ 16.'),
    ('mix_check', 'Meat mix check', '@', 'mixcheck', 'The five shares must add up to 100%.'),
]
for _i in _LINEAS:
    SALIDAS += [('yield%d' % _i, 'Meat line %d: yield (cooked ÷ raw)' % _i, PCT, 'corte%d:J' % _i,
                 'From the "%s" tab.' % H2 if _i == 1 else ''),
                ('raw%d' % _i, 'Meat line %d: raw lb to buy' % _i, DEC1,
                 'if({yield%d}>0,{cooked_lb}*{share%d}/{yield%d},0)' % (_i, _i, _i),
                 'Cooked lb × share ÷ yield.' if _i == 1 else ''),
                ('pieces%d' % _i, 'Meat line %d: pieces (packers, butts, racks…)' % _i, ENTERO, 'piezas%d' % _i,
                 'Raw lb ÷ raw weight per piece on the "%s" tab, rounded up.' % H2 if _i == 1 else ''),
                ('miss%d' % _i, 'Meat line %d: price missing? (1 = yes)' % _i, ENTERO, 'falta%d' % _i, '')]
SALIDAS += [
    ('raw_lb', 'Raw meat to buy, all lines (lb)', DEC1, '+'.join('{raw%d}' % i for i in _LINEAS), ''),
    ('wood_lb', 'Wood or charcoal (lb)', DEC1, '{raw_lb}*{wood_ratio}', ''),
    ('missing', 'Meat lines in use with no price', ENTERO, '+'.join('{miss%d}' % i for i in _LINEAS), ''),
    ('meat_status', 'Meat prices', '@', 'estado',
     'While any price is missing, the food cost uses the package estimate. Type your prices to cost your real menu.'),
    ('c_meat', 'Raw meat cost (only when every price is typed)', MONEDA, 'carne', 'Raw lb × your price per lb, line by line.'),
    ('c_food', 'Food cost (meat, sides, bread and sauces)', MONEDA, 'comida',
     'With your prices: meat + (sides + bread and sauces) × guests. Without them: the package estimate × guests.'),
    ('c_disp', 'Disposables', MONEDA, '{guests}*{disposables}', ''),
    ('c_wood', 'Wood or charcoal', MONEDA, '{wood_lb}*{wood_price}', ''),
    ('c_staff', 'Staff cost (wages + payroll taxes)', MONEDA,
     '(({pit_crew}*{wage_pit}+{servers}*{wage_server})*{event_hours}+{cook_hours}*{wage_pit})*(1+{payroll})', ''),
    ('c_other', 'Rentals and mileage', MONEDA, '{rentals}+{miles}*{per_mile}', ''),
    ('c_total', 'Total direct cost', MONEDA, '{c_food}+{c_disp}+{c_wood}+{c_staff}+{c_other}', ''),
    ('c_guest', 'Direct cost per billable guest', MONEDA, '{c_total}/{billable}', ''),
    ('price', 'PRICE PER GUEST before fees, service charge and tax', MONEDA,
     'roundup({c_total}/(1-{margin}-{commission})/{billable}*2)/2',
     'Rounded up to the next 0.50. Market check: drop-off runs about $15-$30 per person and full-service up to $75 '
     '(q39kc.com); a Los Angeles operator charges $40-58 for parties and $65-85 for weddings (eventective.com).'),
    ('price_lb', 'Equivalent price per lb of cooked meat', MONEDA, '{price}/({oz_guest}/16)',
     'The same price expressed per cooked pound, the way many drop-off menus quote.'),
    ('fee_pit', 'Live pit on site fee', MONEDA, 'if({live_pit}=="Yes",{live_pit_fee},0)', ''),
    ('subtotal', 'Food & service subtotal', MONEDA, '{price}*{billable}+{fee_pit}+{event_fee}',
     'Price per guest × billable guests + live pit fee + event fee.'),
    ('sc_amt', 'Service charge', MONEDA, '{subtotal}*{sc}', ''),
    ('tax_amt', 'Sales tax', MONEDA, '({subtotal}+if({tax_sc}=="Yes",{sc_amt},0))*{tax}',
     'On the subtotal, plus the service charge if your state taxes it.'),
    ('total', 'TOTAL TO QUOTE', MONEDA, '{subtotal}+{sc_amt}+{tax_amt}', ''),
    ('dep_amt', 'Nonrefundable deposit to book the date', MONEDA, '{total}*{deposit}', ''),
    ('balance', 'Balance due before the event', MONEDA, '{total}-{dep_amt}', ''),
    ('comm_amt', 'Marketplace fee', MONEDA, '{subtotal}*{commission}', 'What the marketplace keeps.'),
    ('profit', 'Gross profit on the subtotal', MONEDA, '{subtotal}-{comm_amt}-{c_total}', ''),
    ('margin_calc', 'Gross margin (calculated)', PCT1, 'if({subtotal}==0,0,{profit}/{subtotal})',
     'Calculated from the costs above, never typed. It is at or above the target because the price is rounded up.'),
    ('pit_sug', 'Suggested pit crew for this package (Kit estimate)', ENTERO, 'paq:4', 'From the package table below.'),
    ('srv_sug', 'Suggested servers and attendants for this package (Kit estimate)', ENTERO, 'paq:3',
     'From the package table below.'),
    ('tiers_check', 'Cancellation tiers check', '@', 'tramos', 'Tier days must go down: tier 1 > tier 2 > tier 3.'),
    ('ret1', 'Cancellation in tier 1: amount you keep', MONEDA, '{total}*{r1}', ''),
    ('ret2', 'Cancellation in tier 2: amount you keep', MONEDA, '{total}*{r2}', ''),
    ('ret3', 'Cancellation in tier 3: amount you keep', MONEDA, '{total}*{r3}', ''),
    ('ret4', 'Cancellation in tier 4 or no-show: amount you keep', MONEDA, '{total}*{r4}', ''),
]
S = {k: F_CALC + 1 + i for i, (k, *_r) in enumerate(SALIDAS)}
F_TABLA = F_CALC + 1 + len(SALIDAS) + 1                                # cabecera de la tabla de paquetes
_TP = lambda col: "'%s'!$%s$%d:$%s$%d" % (H1, col, F_TABLA + 1, col, F_TABLA + len(PAQUETES))   # noqa: E731
T_NOMBRES = _TP('A')

# 3 cotizaciones (Kit estimates; D106): cada una a ±10 % del precio de su línea en la mezcla de research §8, SIN precios
# de carne (el coste de comida usa el Kit estimate por invitado del paquete)
OBJETIVO_EJ = OrderedDict([('Corporate drop-off, 40 guests', 22.0), ('Graduation party, staffed buffet, 70 guests', 35.0),
                           ('Wedding, 130 guests, full-service + live pit on site', 52.0)])
_DEF = OrderedDict((k, v) for k, _r, _f2, v, _n in ENTRADAS)


def _ej(**kw):
    e = OrderedDict(_DEF)
    e.update(kw)
    return e


_MEZ_DROP = dict(cut1='Packer brisket (whole, untrimmed)', share1=0.35, cut2='Pork butt (bone-in shoulder)', share2=0.35,
                 cut3='St. Louis ribs', share3=0.0, cut4='Chicken thighs (bone-in)', share4=0.20,
                 cut5='Sausage (smoked links)', share5=0.10)
EJEMPLOS = OrderedDict([
    ('Corporate drop-off, 40 guests', _ej(
        package=DROP, guests=40, minimum=20, pit_crew=1, servers=0, event_hours=4.0, cook_hours=6.0, rentals=0.0,
        miles=20, margin=0.40, **_MEZ_DROP)),
    ('Graduation party, staffed buffet, 70 guests', _ej(
        package=BUFFET, guests=70, minimum=30, pit_crew=1, servers=2, event_hours=6.0, cook_hours=8.0, rentals=0.0,
        miles=30, margin=0.50)),
    ('Wedding, 130 guests, full-service + live pit on site', _ej(
        package=BUFFET, guests=130, minimum=75, live_pit='Yes', pit_crew=2, servers=5, event_hours=8.0, cook_hours=10.0,
        rentals=600.0, miles=60, margin=0.50, sc=0.10)),
])
EJ_NOTAS = OrderedDict([
    ('Corporate drop-off, 40 guests', 'Boxed BBQ lunch for an office, delivered and set (20-guest delivery minimum, as at '
                                      'hutchinsbbq.com). It lands within 10% of the $22 drop-off line in the workbook\'s '
                                      'example mix and inside the $15-$30 per person drop-off range (q39kc.com).'),
    ('Graduation party, staffed buffet, 70 guests', 'Backyard graduation party with a staffed buffet and a carving '
                                                    'station, 30-guest minimum. Within 10% of the $35 party line of the '
                                                    'example mix, below the $40-58 that a Los Angeles operator charges for '
                                                    'parties (eventective.com).'),
    ('Wedding, 130 guests, full-service + live pit on site', 'Full-service wedding with the smoker trailer on site for a '
                                                             '$650 flat fee, rentals passed through and a 10% service '
                                                             'charge; 75-guest minimum. The price per guest, with the '
                                                             'live pit fee spread over the guests, lands within 10% of the '
                                                             '$52 wedding line of the example mix, below the $65-85 of Los '
                                                             'Angeles (eventective.com).'),
])

# --------------------------------------------------------------------------
# 3. Shopping List by Guest Count: parámetros (Kit estimate) y columnas (los invitados de la hoja 1 + T4 del ES)
# --------------------------------------------------------------------------
COMPRA = [
    ('n_sides', 'Number of sides on the menu', ENTERO, 2, 'Kit estimate (3 meats + 2 sides is a common package).'),
    ('side_oz', 'Portion of each side per guest (oz)', DEC1, 4.0, 'Kit estimate.'),
    ('pan_oz', 'Oz of sides per half hotel pan (2.5 in deep)', ENTERO, 128, 'Kit estimate: about a gallon.'),
    ('rolls', 'Rolls or bread slices per guest', DEC1, 1.5, 'Kit estimate.'),
    ('sauce_oz', 'BBQ sauce per guest (oz)', DEC1, 2.0, 'Kit estimate.'),
    ('paper_ft', 'Butcher paper and foil per raw lb (ft of roll)', DEC1, 1.5, 'Kit estimate.'),
    ('lb_pan', 'Cooked meat per full hotel pan (lb)', DEC1, 15.0, 'Kit estimate.'),
    ('pans_box', 'Full hotel pans per insulated food carrier or hot box', ENTERO, 4, 'Kit estimate: check your model.'),
]
REF_INVITADOS = (40, 80, 150, 300)                                    # T4 del ES (calculadora_par_es.py TAM_REF)
C_COLS = 'BCDEF'


# --------------------------------------------------------------------------
# Lógica en Python (la misma que las fórmulas): la usa el autotest como valor esperado
# --------------------------------------------------------------------------
def _ceil(x):
    return math.ceil(round(x, 9))


def _corte(nombre):
    return CORTES[NOMBRES_CORTE.index(nombre)]


def calcular(e, cortes=None):
    cortes = cortes or CORTES
    por_nombre = {c[0]: c for c in cortes}
    r = OrderedDict()
    oz, por_srv, por_pit, kit_food, _n = PAQUETES[e['package']]
    r['billable'] = max(e['guests'], e['minimum'])
    r['oz_guest'] = oz
    r['cooked_lb'] = e['guests'] * oz / 16.0
    tot_share = sum(e['share%d' % i] for i in _LINEAS)
    r['mix_check'] = 'OK' if round(tot_share, 4) == 1 else 'Adjust: the five shares must add up to 100%'
    for i in _LINEAS:
        c = por_nombre[e['cut%d' % i]]
        y, peso = c[7], c[1]
        r['yield%d' % i] = y
        r['raw%d' % i] = r['cooked_lb'] * e['share%d' % i] / y if y > 0 else 0
        r['pieces%d' % i] = _ceil(r['raw%d' % i] / peso) if peso > 0 else 0
        p = e['price%d' % i]
        r['miss%d' % i] = (0 if (isinstance(p, (int, float)) and p > 0) else 1) if e['share%d' % i] > 0 else 0
    r['raw_lb'] = sum(r['raw%d' % i] for i in _LINEAS)
    r['wood_lb'] = r['raw_lb'] * e['wood_ratio']
    r['missing'] = sum(r['miss%d' % i] for i in _LINEAS)
    r['meat_status'] = ('Complete: food cost uses your meat prices' if r['missing'] == 0 else
                        'Missing: food cost uses the package estimate per guest')
    r['c_meat'] = (sum(r['raw%d' % i] * (e['price%d' % i] if isinstance(e['price%d' % i], (int, float)) else 0)
                       for i in _LINEAS) if r['missing'] == 0 else 0)
    r['c_food'] = (r['c_meat'] + e['guests'] * (e['sides'] + e['bread']) if r['missing'] == 0
                   else e['guests'] * kit_food)
    r['c_disp'] = e['guests'] * e['disposables']
    r['c_wood'] = r['wood_lb'] * e['wood_price']
    r['c_staff'] = (((e['pit_crew'] * e['wage_pit'] + e['servers'] * e['wage_server']) * e['event_hours']
                     + e['cook_hours'] * e['wage_pit']) * (1 + e['payroll']))
    r['c_other'] = e['rentals'] + e['miles'] * e['per_mile']
    r['c_total'] = r['c_food'] + r['c_disp'] + r['c_wood'] + r['c_staff'] + r['c_other']
    r['c_guest'] = r['c_total'] / r['billable']
    r['price'] = _ceil(r['c_total'] / (1 - e['margin'] - e['commission']) / r['billable'] * 2) / 2
    r['price_lb'] = r['price'] / (oz / 16.0)
    r['fee_pit'] = e['live_pit_fee'] if e['live_pit'] == 'Yes' else 0
    r['subtotal'] = r['price'] * r['billable'] + r['fee_pit'] + e['event_fee']
    r['sc_amt'] = r['subtotal'] * e['sc']
    r['tax_amt'] = (r['subtotal'] + (r['sc_amt'] if e['tax_sc'] == 'Yes' else 0)) * e['tax']
    r['total'] = r['subtotal'] + r['sc_amt'] + r['tax_amt']
    r['dep_amt'] = r['total'] * e['deposit']
    r['balance'] = r['total'] - r['dep_amt']
    r['comm_amt'] = r['subtotal'] * e['commission']
    r['profit'] = r['subtotal'] - r['comm_amt'] - r['c_total']
    r['margin_calc'] = 0 if r['subtotal'] == 0 else r['profit'] / r['subtotal']
    r['pit_sug'] = _ceil(e['guests'] / por_pit) if por_pit > 0 else 0
    r['srv_sug'] = _ceil(e['guests'] / por_srv) if por_srv > 0 else 0
    r['tiers_check'] = 'OK' if (e['d1'] > e['d2'] > e['d3']) else 'Adjust: tier 1 > tier 2 > tier 3 days'
    for i in range(1, 5):
        r['ret%d' % i] = r['total'] * e['r%d' % i]
    return r


def calcular_yield(par=None):
    """Hoja 2: por corte, °C, raciones por lb cruda, lb cocinadas por pieza, hora de encendido y día."""
    p = OrderedDict((k, v[2]) for k, v in Y_PAR.items())
    if par:
        p.update(par)
    out = []
    for nombre, peso, pit, cal, fc, horas, reposo, rend, fuente, pieza in CORTES:
        x = OrderedDict()
        x['pit_c'] = round((pit - 32) * 5 / 9.0)
        x['cal_c'] = round((cal - 32) * 5 / 9.0)
        x['raciones'] = rend * 16 / p['serving_oz']
        x['coc_pieza'] = peso * rend
        ini = p['service'] - (horas + reposo + p['cushion']) / 24.0
        x['inicio'] = ini - math.floor(ini)
        x['dia'] = 'Event day' if ini >= 0 else ('Day before' if ini >= -1 else '2 days before')
        out.append((nombre, x))
    return out


def calcular_compra(e, c=None, ref=REF_INVITADOS):
    """Hoja 3: para los invitados de la hoja 1 y las 4 columnas de referencia (mismo paquete y misma mezcla)."""
    c = c or OrderedDict((k, v) for k, _r, _f, v, _n in COMPRA)
    oz, por_srv, por_pit, _kf, _n = PAQUETES[e['package']]
    cols = []
    for g in (e['guests'],) + tuple(ref):
        r = OrderedDict()
        r['guests'] = g
        r['cooked'] = g * oz / 16.0
        raw_tot = 0
        for i in _LINEAS:
            ct = _corte(e['cut%d' % i])
            raw = r['cooked'] * e['share%d' % i] / ct[7] if ct[7] > 0 else 0
            r['raw%d' % i] = raw
            r['pz%d' % i] = _ceil(raw / ct[1]) if ct[1] > 0 else 0
            raw_tot += raw
        r['raw_tot'] = raw_tot
        r['sides_gal'] = g * c['n_sides'] * c['side_oz'] / 128.0
        r['half_pans'] = _ceil(g * c['n_sides'] * c['side_oz'] / c['pan_oz'])
        r['rolls'] = _ceil(g * c['rolls'])
        r['sauce_gal'] = g * c['sauce_oz'] / 128.0
        r['wood'] = raw_tot * e['wood_ratio']
        r['paper'] = _ceil(raw_tot * c['paper_ft'])
        r['pans'] = _ceil(r['cooked'] / c['lb_pan'])
        r['boxes'] = _ceil(r['pans'] / c['pans_box'])
        r['pit'] = _ceil(g / por_pit) if por_pit > 0 else 0
        r['srv'] = _ceil(g / por_srv) if por_srv > 0 else 0
        cols.append(r)
    return cols


# --------------------------------------------------------------------------
# Fórmulas
# --------------------------------------------------------------------------
def a_excel(expr, ref):
    """Plantilla Python-ish → fórmula Excel con las celdas de `ref` (clave → coordenada)."""
    f = expr.replace('max(', 'MAX(').replace('roundup(', 'ROUNDUP(').replace('if(', 'IF(').replace('==', '=')
    f = re.sub(r'\{(\w+)\}', lambda m: ref[m.group(1)], f)
    return f.replace('*2)/2', '*2,0)/2')


def _paq(ref, col_idx):
    """Valor de la tabla de paquetes (columna 2 = oz, 3 = invitados por server, 4 = por pit crew, 5 = coste)."""
    return 'INDEX(%s,MATCH(%s,%s,0))' % (_TP('ABCDE'[col_idx - 1]), ref['package'], T_NOMBRES)


def formula_salida(k, expr, ref):
    if expr.startswith('paq:'):
        n = int(expr[4:])
        v = _paq(ref, n)
        if n == 2:
            return '=IFERROR(%s,0)' % v
        return '=IFERROR(IF(%s>0,ROUNDUP(%s/%s,0),0),0)' % (v, ref['guests'], v)
    if expr.startswith('corte'):
        i, col = expr[5:].split(':')
        return '=IFERROR(INDEX(%s,MATCH(%s,%s,0)),0)' % (_ycol(col), ref['cut' + i], T_CORTES)
    if expr.startswith('piezas'):
        i = expr[6:]
        peso = 'INDEX(%s,MATCH(%s,%s,0))' % (_ycol('B'), ref['cut' + i], T_CORTES)
        return '=IFERROR(IF(%s>0,ROUNDUP(%s/%s,0),0),0)' % (peso, ref['raw' + i], peso)
    if expr.startswith('falta'):
        i = expr[5:]
        p = ref['price' + i]
        return '=IF(%s>0,IF(ISNUMBER(%s),IF(%s>0,0,1),1),0)' % (ref['share' + i], p, p)
    if expr == 'mixcheck':
        return '=IF(ROUND(%s,4)=1,"OK","Adjust: the five shares must add up to 100%%")' % '+'.join(
            ref['share%d' % i] for i in _LINEAS)
    if expr == 'estado':
        return ('=IF(%s=0,"Complete: food cost uses your meat prices","Missing: food cost uses the package estimate per '
                'guest")' % ref['missing'])
    if expr == 'carne':
        partes = ['IF(ISNUMBER(%s),%s*%s,0)' % (ref['price%d' % i], ref['raw%d' % i], ref['price%d' % i]) for i in _LINEAS]
        return '=IF(%s=0,%s,0)' % (ref['missing'], '+'.join(partes))
    if expr == 'comida':
        return '=IF(%s=0,%s+%s*(%s+%s),%s*IFERROR(%s,0))' % (ref['missing'], ref['c_meat'], ref['guests'], ref['sides'],
                                                            ref['bread'], ref['guests'], _paq(ref, 5))
    if expr == 'tramos':
        return '=IF(AND(%s>%s,%s>%s),"OK","Adjust: tier 1 > tier 2 > tier 3 days")' % (ref['d1'], ref['d2'], ref['d2'],
                                                                                      ref['d3'])
    return '=' + a_excel(expr, ref)


# --------------------------------------------------------------------------
# Construcción
# --------------------------------------------------------------------------
def _f(ws, coord, v, b=False, sz=11, i=False, fmt=None, wrap=False, fill=None, locked=True, color=None):
    c = ws[coord]
    c.value = v
    c.font = Font(name=F, size=sz, bold=b, italic=i, color=color)
    if fmt:
        c.number_format = fmt
    c.alignment = Alignment(wrap_text=wrap, vertical='top')
    if fill:
        c.fill = fill
    c.protection = Protection(locked=locked)
    return c


def _imprimir(ws, horizontal=False):
    ws.page_setup.paperSize = M.PAPEL_EN                                         # 1 = US Letter
    ws.page_setup.orientation = 'landscape' if horizontal else 'portrait'
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.oddFooter.center.text = M.FOOTER_EN
    ws.print_options.horizontalCentered = True
    ws.protection.sheet = True                                                   # sin contraseña (D20)
    ws.sheet_view.showGridLines = False


def _dv(ws, tipo, celdas, f1=None, f2=None, op=None, titulo='Invalid value', msg='', blanco=False):
    dv = DataValidation(type=tipo, formula1=f1, formula2=f2, operator=op, allow_blank=blanco, showErrorMessage=True,
                        errorTitle=titulo, error=msg)
    ws.add_data_validation(dv)
    for c in celdas:
        dv.add(c)


def _cabecera(ws, titulo, intro, ancho_merge, alto=48):
    _f(ws, 'A1', titulo, b=True, sz=14)
    _f(ws, 'A2', intro, wrap=True)
    ws.merge_cells('A2:%s2' % ancho_merge)
    ws.row_dimensions[2].height = alto


def _fmt_entrada(fmt):
    return fmt if fmt != '@' else None


def hoja_calculadora(wb):
    ws = wb.active
    ws.title = H1
    for col, w in (('A', 62), ('B', 34), ('C', 80), ('D', 16), ('E', 22), ('F', 70)):
        ws.column_dimensions[col].width = w
    _cabecera(ws, 'BBQ CATERING PRICING CALCULATOR — POUNDS OF MEAT PER GUEST TO QUOTE',
              'Type in the GREEN cells; everything else is calculated. Pick the package and the five meat lines, type your '
              'raw meat prices per lb (left blank on purpose), your staff and your costs: the tab turns them into pounds '
              'to buy, a price per guest and a full quote. Amounts are in US dollars (or your own currency) and exclude '
              'sales tax unless the row says otherwise. Every default value is a kit estimate: replace it with your own '
              'numbers.', 'C')
    for col, t in zip('ABC', ('INPUTS', 'Value', 'Notes')):
        _f(ws, '%s4' % col, t, b=True, fill=GRIS)
    for k, rot, fmt, val, nota in ENTRADAS:
        f = K[k]
        _f(ws, 'A%d' % f, rot)
        _f(ws, 'B%d' % f, val, fmt=_fmt_entrada(fmt), fill=VERDE, locked=False, wrap=k == 'package' or k.startswith('cut'))
        if nota:
            _f(ws, 'C%d' % f, nota, wrap=True)
        if len(nota) > 90 or k == 'package':
            ws.row_dimensions[f].height = 15 * math.ceil(len(nota) / 90.0) + 2
    for col, t in zip('ABC', ('CALCULATION', 'Value', 'Notes')):
        _f(ws, '%s%d' % (col, F_CALC), t, b=True, fill=GRIS)
    ref = {k: '$B$%d' % f for k, f in K.items()}
    ref.update({k: '$B$%d' % f for k, f in S.items()})
    for k, rot, fmt, expr, nota in SALIDAS:
        f = S[k]
        neg = k in ('price', 'total', 'margin_calc', 'meat_status')
        _f(ws, 'A%d' % f, rot, b=neg)
        _f(ws, 'B%d' % f, formula_salida(k, expr, ref), fmt=_fmt_entrada(fmt), b=neg, wrap=fmt == '@')
        if nota:
            _f(ws, 'C%d' % f, nota, wrap=True)
            if len(nota) > 90:
                ws.row_dimensions[f].height = 15 * math.ceil(len(nota) / 90.0) + 2
        elif fmt == '@':
            ws.row_dimensions[f].height = 32
    cab = ('PACKAGES (Kit estimate — edit the numbers)', 'Cooked meat per guest (oz)', 'Guests per server or attendant '
           '(0 = none)', 'Guests per pit crew member', 'Food cost per guest while meat prices are blank', 'Notes')
    for col, t in zip('ABCDEF', cab):
        _f(ws, '%s%d' % (col, F_TABLA), t, b=True, fill=GRIS, wrap=True)
    ws.row_dimensions[F_TABLA].height = 45
    for i, (paq, (oz, srv, pit, food, nota)) in enumerate(PAQUETES.items(), 1):
        f = F_TABLA + i
        _f(ws, 'A%d' % f, paq, wrap=True)
        _f(ws, 'B%d' % f, oz, fmt=DEC1, fill=VERDE, locked=False)
        _f(ws, 'C%d' % f, srv, fmt=ENTERO, fill=VERDE, locked=False)
        _f(ws, 'D%d' % f, pit, fmt=ENTERO, fill=VERDE, locked=False)
        _f(ws, 'E%d' % f, food, fmt=MONEDA, fill=VERDE, locked=False)
        _f(ws, 'F%d' % f, nota, wrap=True)
        ws.row_dimensions[f].height = 48
    f = F_TABLA + len(PAQUETES) + 1
    _f(ws, 'A%d' % f, 'The food cost per guest in this table (Kit estimates: $6.50 drop-off, $8 buffet, $17 live-fire) '
                      'is only a placeholder while your meat prices are blank. It is not a meat price: as soon as every '
                      'meat line in use has a price, the calculator costs your real pounds instead.', wrap=True, i=True)
    ws.merge_cells('A%d:F%d' % (f, f))
    ws.row_dimensions[f].height = 32
    # validaciones (D105)
    _dv(ws, 'list', [ref['package'].replace('$', '')], f1=T_NOMBRES.replace("'%s'!" % H1, ''),
        titulo='Invalid package', msg='Choose a package from the list.')
    _dv(ws, 'list', [ref['cut%d' % i].replace('$', '') for i in _LINEAS], f1=T_CORTES, titulo='Invalid cut',
        msg='Choose a cut from the list on the "%s" tab.' % H2)
    _dv(ws, 'list', [ref[k].replace('$', '') for k in ('tax_sc', 'live_pit')], f1='"Yes,No"', titulo='Invalid entry',
        msg='Choose Yes or No.')
    _dv(ws, 'whole', ['B%d' % K['guests']], f1='1', f2='100000', op='between', titulo='Invalid count',
        msg='Enter a whole number of 1 or more.')
    _dv(ws, 'whole', ['B%d' % K[k] for k in ('minimum', 'pit_crew', 'servers', 'miles')], f1='0', f2='100000',
        op='between', titulo='Invalid count', msg='Enter a whole number of 0 or more.')
    _dv(ws, 'whole', ['B%d' % K[k] for k in ('d1', 'd2', 'd3')], f1='1', f2='730', op='between', titulo='Invalid days',
        msg='Enter a whole number of days between 1 and 730.')
    _dv(ws, 'decimal', ['B%d' % K['price%d' % i] for i in _LINEAS], f1='0', op='greaterThanOrEqual',
        titulo='Invalid price', msg='Type your raw price per lb (a number of 0 or more), or leave it blank.', blanco=True)
    _dv(ws, 'decimal', ['B%d' % K[k] for k in ('sides', 'bread', 'disposables', 'wood_ratio', 'wood_price', 'live_pit_fee',
                                              'event_hours', 'cook_hours', 'wage_pit', 'wage_server', 'rentals',
                                              'per_mile', 'event_fee')],
        f1='0', op='greaterThanOrEqual', titulo='Invalid amount', msg='Enter a number greater than or equal to 0.')
    _dv(ws, 'decimal', ['B%d' % K[k] for k in ['share%d' % i for i in _LINEAS] + ['payroll', 'sc', 'tax', 'deposit', 'r1',
                                                                                 'r2', 'r3', 'r4']],
        f1='0', f2='1', op='between', titulo='Invalid percentage', msg='Enter a percentage between 0 and 1 (0.20 = 20%).')
    _dv(ws, 'decimal', ['B%d' % K['commission']], f1='0', f2='0.3', op='between', titulo='Invalid percentage',
        msg='Enter a fee between 0 and 30% (0.15 = 15%).')
    _dv(ws, 'decimal', ['B%d' % K['margin']], f1='0', f2='0.65', op='between', titulo='Invalid percentage',
        msg='Enter a target margin between 0 and 65% (0.45 = 45%).')
    _dv(ws, 'decimal', ['B%d' % (F_TABLA + i) for i in range(1, len(PAQUETES) + 1)], f1='0.5', f2='32', op='between',
        titulo='Invalid amount', msg='Enter the cooked oz of meat per guest, between 0.5 and 32.')
    _dv(ws, 'whole', ['%s%d' % (c, F_TABLA + i) for i in range(1, len(PAQUETES) + 1) for c in 'CD'], f1='0',
        f2='10000', op='between', titulo='Invalid count', msg='Enter a whole number of guests, 0 or more.')
    _dv(ws, 'decimal', ['E%d' % (F_TABLA + i) for i in range(1, len(PAQUETES) + 1)], f1='0', op='greaterThanOrEqual',
        titulo='Invalid amount', msg='Enter a number greater than or equal to 0.')
    ws.freeze_panes = 'A5'
    _imprimir(ws)
    return ws


def hoja_yield(wb):
    ws = wb.create_sheet(H2)
    anchos = (('A', 36), ('B', 11), ('C', 10), ('D', 8), ('E', 11), ('F', 8), ('G', 34), ('H', 9), ('I', 9), ('J', 9),
              ('K', 26), ('L', 11), ('M', 11), ('N', 12), ('O', 13))
    for col, w in anchos:
        ws.column_dimensions[col].width = w
    _cabecera(ws, 'MEAT YIELD & SMOKE SCHEDULE — 13 CUTS',
              'Raw weight, pit temperature, cook and rest times and yield for 13 cuts (green cells: kit estimates, edit '
              'them with your own test cooks). The QUALITY target (when the meat is done the way guests expect) and the '
              'FDA Food Code MINIMUM (the legal safety floor, 2022 Food Code 3-401.11) are different numbers and sit in '
              'separate columns. Type your service time below: the tab tells you when to start each cook.', 'O', alto=60)
    _f(ws, 'A4', 'SCHEDULE INPUTS', b=True, fill=GRIS)
    for k, (rot, fmt, val, nota) in Y_PAR.items():
        f = Y_PAR_FILA[k]
        _f(ws, 'A%d' % f, rot)
        _f(ws, 'B%d' % f, val, fmt=fmt, fill=VERDE, locked=False)
        _f(ws, 'C%d' % f, nota)
    cab = ('Cut', 'Raw weight per piece (lb)', 'Pit or grill temp (°F)', '(°C)', 'Quality target (°F, Kit estimate)',
           '(°C)', 'FDA Food Code minimum (safety)', 'Cook hours', 'Rest and hold hours', 'Yield (cooked ÷ raw)',
           'Yield source', 'Servings per raw lb', 'Cooked lb per piece', 'Start the pit by', 'Day')
    fila = Y_F0 - 1
    for col, t in zip('ABCDEFGHIJKLMNO', cab):
        _f(ws, '%s%d' % (col, fila), t, b=True, fill=GRIS, wrap=True)
    ws.row_dimensions[fila].height = 45
    g = {k: '$B$%d' % f for k, f in Y_PAR_FILA.items()}
    verdes = []
    for i, (nombre, peso, pit, cal, fc, horas, reposo, rend, fuente, _pz) in enumerate(CORTES):
        r = Y_F0 + i
        _f(ws, 'A%d' % r, nombre)
        _f(ws, 'B%d' % r, peso, fmt=DEC2, fill=VERDE, locked=False)
        _f(ws, 'C%d' % r, pit, fmt=ENTERO, fill=VERDE, locked=False)
        _f(ws, 'D%d' % r, '=ROUND((C%d-32)*5/9,0)' % r, fmt=ENTERO)
        _f(ws, 'E%d' % r, cal, fmt=ENTERO, fill=VERDE, locked=False)
        _f(ws, 'F%d' % r, '=ROUND((E%d-32)*5/9,0)' % r, fmt=ENTERO)
        _f(ws, 'G%d' % r, fc, wrap=True)
        _f(ws, 'H%d' % r, horas, fmt=DEC2, fill=VERDE, locked=False)
        _f(ws, 'I%d' % r, reposo, fmt=DEC2, fill=VERDE, locked=False)
        _f(ws, 'J%d' % r, rend, fmt=PCT, fill=VERDE, locked=False)
        _f(ws, 'K%d' % r, fuente, wrap=True)
        _f(ws, 'L%d' % r, '=IF(%s>0,J%d*16/%s,0)' % (g['serving_oz'], r, g['serving_oz']), fmt=DEC1)
        _f(ws, 'M%d' % r, '=B%d*J%d' % (r, r), fmt=DEC1)
        ini = '(%s-(H%d+I%d+%s)/24)' % (g['service'], r, r, g['cushion'])
        _f(ws, 'N%d' % r, '=%s-INT(%s)' % (ini, ini), fmt=HORA)
        _f(ws, 'O%d' % r, '=IF(%s>=0,"Event day",IF(%s>=-1,"Day before","2 days before"))' % (ini, ini))
        ws.row_dimensions[r].height = 30
        verdes += ['%s%d' % (c, r) for c in 'BCEHI']
    _dv(ws, 'decimal', verdes, f1='0', op='greaterThanOrEqual', titulo='Invalid amount',
        msg='Enter a number greater than or equal to 0.')
    _dv(ws, 'decimal', ['J%d' % (Y_F0 + i) for i in range(len(CORTES))], f1='0.05', f2='1', op='between',
        titulo='Invalid yield', msg='Enter a yield between 0.05 and 1 (0.50 = 50%).')
    _dv(ws, 'time', [g['service'].replace('$', '')], f1='0', f2='0.999988', op='between', titulo='Invalid time',
        msg='Type a time of day, such as 6:00 PM.')
    _dv(ws, 'decimal', [g['cushion'].replace('$', '')], f1='0', f2='12', op='between', titulo='Invalid hours',
        msg='Enter the cushion in hours, between 0 and 12.')
    _dv(ws, 'decimal', [g['serving_oz'].replace('$', '')], f1='0.5', f2='32', op='between', titulo='Invalid amount',
        msg='Enter the cooked oz per serving, between 0.5 and 32.')
    fila = Y_F0 + len(CORTES) + 1
    for t in ('Quality target = the internal temperature at which the cut is done the way guests expect (Kit estimate: '
              'probe-tender brisket around 203 °F). Food Code minimum = the legal safety floor. Pull a cut when it '
              'reaches BOTH; never write one in place of the other.',
              'Injecting a brisket or a pork butt makes it non-intact: its minimum becomes 155 °F (68 °C) for 17 seconds. '
              'Serving whole-muscle beef below its minimum (a medium-rare tri-tip, picanha or tomahawk): check your county '
              'rules before you put it on the menu.',
              'Hot holding: keep cooked meat at 135 °F (57 °C) or above (3-501.16), in insulated food carriers or hot '
              'boxes, and probe it before you serve.',
              'Cooling leftovers: to 70 °F (21 °C) within 2 hours and to 41 °F (5 °C) within a total of 6 hours '
              '(3-501.14); slice or portion brisket so it cools fast. Reheating cooked meat for hot holding: to 165 °F '
              '(74 °C) for 15 seconds (3-403.11).',
              'Yield: brisket and pork butt end up at about 50% of their raw weight (fieryfoodscentral.com, Dr. BBQ); '
              'every other yield is a kit estimate. Weigh your own raw and cooked pounds on your test cook and type '
              'them here: the "%s" and "%s" tabs follow.' % (H1, H3),
              'Whole hog: buy it from a USDA- or state-inspected plant; custom-exempt meat is marked "not for sale" '
              '(nichemeatprocessing.org).',
              'Start the pit by = service time − cook hours − rest and hold hours − safety cushion. Long cooks start the '
              'day before: plan the overnight shift of your pit cook.'):
        _f(ws, 'A%d' % fila, t, wrap=True)
        ws.merge_cells('A%d:O%d' % (fila, fila))
        ws.row_dimensions[fila].height = 32
        fila += 1
    ws.freeze_panes = 'B%d' % Y_F0
    _imprimir(ws, horizontal=True)
    return ws


C_FILAS = None                                                         # se rellena en hoja_compra (fila de cada clave)


def hoja_compra(wb):
    ws = wb.create_sheet(H3)
    for col, w in (('A', 50), ('B', 16), ('C', 14), ('D', 14), ('E', 14), ('F', 14), ('G', 56)):
        ws.column_dimensions[col].width = w
    _cabecera(ws, 'SHOPPING LIST BY GUEST COUNT — WHAT TO BUY AND BRING',
              'It reads the package, the meat lines (cut and share) and the wood ratio from the "%s" tab and the yields '
              'and piece weights from the "%s" tab. Column B is your event; columns C-F are reference sizes (40, 80, 150 '
              'and 300 guests: edit them). Green cells are kit estimates: edit them.' % (H1, H2), 'G')
    g1 = "'%s'!" % H1
    R = {}
    fila = 4
    _f(ws, 'A%d' % fila, 'ASSUMPTIONS (Kit estimate — edit)', b=True, fill=GRIS)
    for col, t in zip('BG', ('Value', 'Notes')):
        _f(ws, '%s%d' % (col, fila), t, b=True, fill=GRIS)
    for k, rot, fmt, val, nota in COMPRA:
        fila += 1
        _f(ws, 'A%d' % fila, rot)
        _f(ws, 'B%d' % fila, val, fmt=fmt, fill=VERDE, locked=False)
        if nota:
            _f(ws, 'G%d' % fila, nota, wrap=True)
        R[k] = '$B$%d' % fila
    fila += 2
    f_cab = fila
    _f(ws, 'A%d' % fila, 'QUANTITIES', b=True, fill=GRIS)
    _f(ws, 'B%d' % fila, 'Your event', b=True, fill=GRIS)
    for j in range(4):
        _f(ws, '%s%d' % (C_COLS[j + 1], fila), 'Reference', b=True, fill=GRIS)
    _f(ws, 'G%d' % fila, 'Notes', b=True, fill=GRIS)
    oz = 'IFERROR(INDEX(%s,MATCH(%s$B$%d,%s,0)),0)' % (_TP('B'), g1, K['package'], T_NOMBRES)
    srv = 'IFERROR(INDEX(%s,MATCH(%s$B$%d,%s,0)),0)' % (_TP('C'), g1, K['package'], T_NOMBRES)
    pit = 'IFERROR(INDEX(%s,MATCH(%s$B$%d,%s,0)),0)' % (_TP('D'), g1, K['package'], T_NOMBRES)
    filas = OrderedDict()

    def linea(clave, rot, fx, fmt, nota=''):
        nonlocal fila
        fila += 1
        _f(ws, 'A%d' % fila, rot)
        for j, col in enumerate(C_COLS):
            _f(ws, '%s%d' % (col, fila), fx(col), fmt=fmt)
        if nota:
            _f(ws, 'G%d' % fila, nota, wrap=True)
        filas[clave] = fila

    # invitados: B = hoja 1; C-F = verdes editables
    fila += 1
    _f(ws, 'A%d' % fila, 'Guests', b=True)
    _f(ws, 'B%d' % fila, '=%s$B$%d' % (g1, K['guests']), fmt=ENTERO, b=True)
    for j, n in enumerate(REF_INVITADOS):
        _f(ws, '%s%d' % (C_COLS[j + 1], fila), n, fmt=ENTERO, b=True, fill=VERDE, locked=False)
    _f(ws, 'G%d' % fila, 'Column B comes from the "%s" tab.' % H1, wrap=True)
    filas['guests'] = fila
    gr = lambda col: '%s$%d' % (col, filas['guests'])                           # noqa: E731
    linea('cooked', 'Cooked meat to serve (lb)', lambda c: '=%s*%s/16' % (gr(c), oz), DEC1, 'Guests × oz per guest of '
          'the package ÷ 16.')
    for i in _LINEAS:
        corte = '%s$B$%d' % (g1, K['cut%d' % i])
        share = '%s$B$%d' % (g1, K['share%d' % i])
        y = 'INDEX(%s,MATCH(%s,%s,0))' % (_ycol('J'), corte, T_CORTES)
        peso = 'INDEX(%s,MATCH(%s,%s,0))' % (_ycol('B'), corte, T_CORTES)
        fila += 1
        _f(ws, 'A%d' % fila, '="Raw lb: "&%s' % corte)
        ck = filas['cooked']
        for col in C_COLS:
            _f(ws, '%s%d' % (col, fila), '=IFERROR(IF(%s>0,%s$%d*%s/%s,0),0)' % (y, col, ck, share, y), fmt=DEC1)
        filas['raw%d' % i] = fila
        fila += 1
        _f(ws, 'A%d' % fila, '="   Pieces: "&%s' % corte)
        for col in C_COLS:
            _f(ws, '%s%d' % (col, fila), '=IFERROR(IF(%s>0,ROUNDUP(%s$%d/%s,0),0),0)' % (peso, col, fila - 1, peso),
               fmt=ENTERO)
        filas['pz%d' % i] = fila
    linea('raw_tot', 'Raw meat, all lines (lb)', lambda c: '=' + '+'.join('%s$%d' % (c, filas['raw%d' % i])
                                                                          for i in _LINEAS), DEC1, '')
    linea('sides_gal', 'Sides (gallons)', lambda c: '=%s*%s*%s/128' % (gr(c), R['n_sides'], R['side_oz']), DEC1,
          '1 gallon = 128 oz (3.8 liters).')
    linea('half_pans', 'Sides (half hotel pans)', lambda c: '=ROUNDUP(%s*%s*%s/%s,0)' % (gr(c), R['n_sides'],
                                                                                         R['side_oz'], R['pan_oz']),
          ENTERO, '')
    linea('rolls', 'Rolls or bread slices', lambda c: '=ROUNDUP(%s*%s,0)' % (gr(c), R['rolls']), ENTERO, '')
    linea('sauce_gal', 'BBQ sauce (gallons)', lambda c: '=%s*%s/128' % (gr(c), R['sauce_oz']), DEC2, '')
    linea('wood', 'Wood or charcoal (lb)', lambda c: '=%s$%d*%s$B$%d' % (c, filas['raw_tot'], g1, K['wood_ratio']), DEC1,
          'Raw lb × the wood ratio of the "%s" tab.' % H1)
    linea('paper', 'Butcher paper and foil (ft of roll)', lambda c: '=ROUNDUP(%s$%d*%s,0)' % (c, filas['raw_tot'],
                                                                                              R['paper_ft']), ENTERO, '')
    linea('pans', 'Full hotel pans for the meat', lambda c: '=ROUNDUP(%s$%d/%s,0)' % (c, filas['cooked'], R['lb_pan']),
          ENTERO, '')
    linea('boxes', 'Insulated food carriers or hot boxes', lambda c: '=ROUNDUP(%s$%d/%s,0)' % (c, filas['pans'],
                                                                                               R['pans_box']), ENTERO,
          'Keep the meat at 135 °F (57 °C) or above until you serve it.')
    linea('pit', 'Pit crew (suggested)', lambda c: '=IF(%s>0,ROUNDUP(%s/%s,0),0)' % (pit, gr(c), pit), ENTERO,
          'From the package table of the "%s" tab.' % H1)
    linea('srv', 'Servers and attendants (suggested)', lambda c: '=IF(%s>0,ROUNDUP(%s/%s,0),0)' % (srv, gr(c), srv),
          ENTERO, '0 for drop-off.')
    _dv(ws, 'whole', ['%s%d' % (C_COLS[j + 1], filas['guests']) for j in range(4)], f1='1', f2='100000', op='between',
        titulo='Invalid count', msg='Enter a whole number of guests, 1 or more.')
    _dv(ws, 'whole', [R[k].replace('$', '') for k in ('n_sides', 'pan_oz', 'pans_box')], f1='0', f2='10000',
        op='between', titulo='Invalid count', msg='Enter a whole number of 0 or more.')
    _dv(ws, 'decimal', [R[k].replace('$', '') for k in ('side_oz', 'rolls', 'sauce_oz', 'paper_ft', 'lb_pan')], f1='0',
        op='greaterThanOrEqual', titulo='Invalid amount', msg='Enter a number greater than or equal to 0.')
    ws.freeze_panes = 'B%d' % (f_cab + 1)
    _imprimir(ws, horizontal=True)
    return ws, filas


C_CLAVES = ('guests', 'cooked', 'raw1', 'pz1', 'raw2', 'pz2', 'raw3', 'pz3', 'raw4', 'pz4', 'raw5', 'pz5', 'raw_tot',
            'sides_gal', 'half_pans', 'rolls', 'sauce_gal', 'wood', 'paper', 'pans', 'boxes', 'pit', 'srv')


def hoja_ejemplos(wb):
    ws = wb.create_sheet(H4)
    ws.column_dimensions['A'].width = 62
    for col in 'BCD':
        ws.column_dimensions[col].width = 30
    _cabecera(ws, '3 EXAMPLE QUOTES — calculated with the same formulas as the "%s" tab' % H1,
              'Kit estimates for illustration. No meat price is typed here: until you type yours, the food cost uses the '
              'package estimate per guest, as in the "%s" tab. Each quote lands within 10%% of the price of its line in '
              'the workbook\'s example mix ($22 drop-off, $35 party, $52 wedding). Change the inputs in "%s" to price '
              'your own events.' % (H1, H1), 'D')
    _f(ws, 'A4', 'INPUTS', b=True, fill=GRIS)
    for j, nombre in enumerate(EJEMPLOS):
        _f(ws, '%s4' % 'BCD'[j], nombre, b=True, fill=GRIS, wrap=True)
    ws.row_dimensions[4].height = 45
    for k, rot, fmt, val, nota in ENTRADAS:
        _f(ws, 'A%d' % K[k], rot)
        for j, e in enumerate(EJEMPLOS.values()):
            if e[k] is not None:
                _f(ws, '%s%d' % ('BCD'[j], K[k]), e[k], fmt=_fmt_entrada(fmt), wrap=k == 'package' or k.startswith('cut'))
            else:
                _f(ws, '%s%d' % ('BCD'[j], K[k]), None, fmt=_fmt_entrada(fmt))
        if k == 'package' or k.startswith('cut'):
            ws.row_dimensions[K[k]].height = 32
    _f(ws, 'A%d' % F_CALC, 'RESULT', b=True, fill=GRIS)
    for j in range(3):
        _f(ws, '%s%d' % ('BCD'[j], F_CALC), list(EJEMPLOS)[j], b=True, fill=GRIS, wrap=True)
    ws.row_dimensions[F_CALC].height = 45
    for k, rot, fmt, expr, nota in SALIDAS:
        neg = k in ('price', 'total', 'margin_calc', 'meat_status')
        _f(ws, 'A%d' % S[k], rot, b=neg)
        for j in range(3):
            col = 'BCD'[j]
            ref = {kk: '$%s$%d' % (col, f) for kk, f in K.items()}
            ref.update({kk: '$%s$%d' % (col, f) for kk, f in S.items()})
            _f(ws, '%s%d' % (col, S[k]), formula_salida(k, expr, ref), fmt=_fmt_entrada(fmt), b=neg, wrap=fmt == '@')
        if fmt == '@':
            ws.row_dimensions[S[k]].height = 32
    f = S[SALIDAS[-1][0]] + 2
    _f(ws, 'A%d' % f, 'Price per guest with the live pit and event fees spread over the billable guests', b=True)
    for j in range(3):
        col = 'BCD'[j]
        _f(ws, '%s%d' % (col, f), '=$%s$%d/$%s$%d' % (col, S['subtotal'], col, S['billable']), fmt=MONEDA, b=True)
    f_equiv = f
    f += 2
    _f(ws, 'A%d' % f, 'WHAT EACH EXAMPLE SHOWS', b=True, fill=GRIS)
    for j, (nombre, nota) in enumerate(EJ_NOTAS.items()):
        _f(ws, 'A%d' % (f + 1 + j), nombre, b=True, wrap=True)
        _f(ws, 'B%d' % (f + 1 + j), nota, wrap=True)
        ws.merge_cells('B%d:D%d' % (f + 1 + j, f + 1 + j))
        ws.row_dimensions[f + 1 + j].height = 62
    ws.freeze_panes = 'B5'
    _imprimir(ws, horizontal=True)
    return ws, f_equiv


INSTRUCCIONES = [
    ('HOW TO USE THE BBQ CATERING PRICING CALCULATOR', True),
    ('This workbook turns pounds of meat into a price per guest and a full quote, tells you when to start each cook and '
     'builds the shopping list. You type only in the GREEN cells; the gross margin is calculated, never typed.', False),
    ('1. "%s": choose the package (drop-off, full-service BBQ buffet or live-fire premium), the guests and your guest '
     'minimum. The package sets the cooked meat per guest: 6 oz for drop-off and 8 oz (half a pound) for the buffet and '
     'live-fire, inside the "one-third to one-half pound of meat per guest" of hutchinsbbq.com and the 6-8 oz of '
     'sonnybryans.com; 1/4 to 1/3 lb of cooked meat (Dr. BBQ, fieryfoodscentral.com) is the floor.' % H1, False),
    ('2. Pick the cut and the share of each of the five meat lines (they must add up to 100%) and type your raw price '
     'per lb for every line you use. The prices are blank on purpose: this kit does not guess 2026 meat prices. While '
     'any of them is missing, the food cost uses the package estimate per guest and the "Meat prices" row says so.',
     False),
    ('3. Enter sides, bread and sauces, disposables, wood, the live pit on site switch and its fee, your staff (event '
     'day and overnight cook hours), rentals and mileage. Then the marketplace fee, your target margin, the service '
     'charge, the sales tax rate and the switch "Does your state tax the service charge?". The price per guest is the '
     'direct cost divided by (1 - target margin - marketplace fee) and by the billable guests, rounded up to the next '
     '0.50; the live pit fee and any event fee go on top. The tab also shows the same price per cooked lb.', False),
    ('4. Deposit and cancellation: a 50% nonrefundable deposit books the date (the operators read for this kit ask for '
     'it: eventective.com, hutchinsbbq.com) and four cancellation tiers with no gaps tell you what you keep. Put the same '
     'tiers in your contract, with a force majeure clause for a county burn ban (propane backup or a new date).', False),
    ('5. "%s": the raw weight, pit temperature, quality target, Food Code minimum, cook and rest hours and yield of 13 '
     'cuts, and the time to start each cook for your service time. The meat lines of the first tab read their yield and '
     'piece weight here: replace the kit estimates with the numbers from your own test cook.' % H2, False),
    ('6. "%s": pounds and pieces per meat line, sides, rolls, sauce, wood, butcher paper, hotel pans, hot boxes and the '
     'suggested crew for your event and for 40, 80, 150 and 300 guests.' % H3, False),
    ('7. "%s": three complete quotes calculated with the same formulas. To move from these per-package prices to the '
     'average price per guest of the financial projections, weight each package price by its guests in the year.' % H4,
     False),
    ('FOOD SAFETY (FDA FOOD CODE 2022)', True),
    ('Cook to the minimum: 145 °F (63 °C) for 15 seconds for intact meat; 145 °F for 4 minutes for whole roasts; 155 °F '
     '(68 °C) for 17 seconds for injected, tenderized or ground meat, sausage included; 165 °F (74 °C) for poultry '
     '(3-401.11). Hold hot at 135 °F (57 °C) or above (3-501.16). Cool to 70 °F (21 °C) within 2 hours and to 41 °F '
     '(5 °C) within a total of 6 hours (3-501.14). Reheat for hot holding to 165 °F (3-403.11). The quality target of '
     'each cut is a different number: see the "%s" tab.' % H2, False),
    ('SERVICE CHARGE, GRATUITY AND SALES TAX (US)', True),
    ('A mandatory service charge is not a tip: the part you pay out to the staff is wages, with payroll taxes (IRS Rev. '
     'Rul. 2012-18), and the part you keep is revenue. A voluntary gratuity is the client\'s decision and belongs to the '
     'staff: this calculator keeps it out of the price and the total.', False),
    ('Sales tax: many states tax the whole catering invoice (food, staff, rentals and fees); whether the service charge '
     'is taxable depends on the state (California taxes a mandatory one: cdtfa.ca.gov). Set your rate and the switch.',
     False),
    ('LIVE PIT ON SITE', True),
    ('Towing the smoker trailer to the event and cooking there is a premium you charge for, not the base of the business. '
     'Ask your county first: cooking at the event can make the trailer a mobile food unit that must operate from your '
     'commissary, and the fire marshal sets the rules for solid fuel, propane and tents.', False),
    ('UK: if you are VAT-registered, enter 0.20 as the tax rate: VAT applies to catering for events and functions.', False),
    ('KIT ESTIMATES', True),
    ('Every default value, every yield except brisket and pork butt, the package food cost per guest and the three '
     'example quotes are kit estimates for illustration, not market prices. Replace them with your own wages, supplier '
     'prices, yields and tax rate before you quote.', False),
]


def hoja_instrucciones(wb):
    ws = wb.create_sheet(H5)
    ws.column_dimensions['A'].width = 110
    f = 1
    for t, b in INSTRUCCIONES:
        _f(ws, 'A%d' % f, t, b=b, sz=13 if f == 1 else 11, wrap=True)
        if not b:
            ws.row_dimensions[f].height = max(16, 15 * math.ceil(len(t) / 120.0))
        f += 2 if b and f > 1 else 1
    f += 1
    for t in ('To edit a cell that is not green: Review → Unprotect Sheet (there is no password).',
              'More templates and kits in the catalog at ' + M.URL_TIENDA, M.FIRMA_EN, M.version_en('bbq'),
              M.AVISO_EN):
        _f(ws, 'A%d' % f, t, wrap=True, i=t == M.AVISO_EN)
        f += 1
    _imprimir(ws)
    return ws


def docprops(wb):
    p = M.PLANES['bbq']
    pr = wb.properties
    pr.title = '%s · %s' % (TITULO, p['producto'])
    pr.subject = '%s · v%s' % (p['producto'], M.VERSION)
    pr.keywords = p['keywords']
    pr.description = p['url']
    pr.category = 'AI Chef Pro · Digital products'
    pr.creator = 'AI Chef Pro'
    pr.lastModifiedBy = 'AI Chef Pro'


def generar(destino):
    wb = openpyxl.Workbook()
    hoja_calculadora(wb)
    hoja_yield(wb)
    hoja_compra(wb)
    hoja_ejemplos(wb)
    hoja_instrucciones(wb)
    docprops(wb)
    os.makedirs(os.path.dirname(destino), exist_ok=True)
    wb.save(destino)
    sys.path.insert(0, PD)
    import inject_cache as IC                                                    # noqa: E402
    nf, inj, fail = IC.inject(destino)
    print('%s: formulas=%d cache_inyectado=%d fallos_pycel=%d' % (os.path.relpath(destino, REPO), nf, inj, fail))
    return fail


# --------------------------------------------------------------------------
# Autotest (D106): pycel = caché = Python en las 4 pestañas de cálculo + el paquete live-fire, el interruptor del live
# pit y los precios de carne tecleados probados en la hoja principal; mezcla = 100 %; ejemplos a ±10 %
# --------------------------------------------------------------------------
def _cmp(err, etiqueta, got, esperado, tol=0.005):
    if isinstance(esperado, str):
        if got != esperado:
            err.append('%s: %r ≠ esperado %r' % (etiqueta, got, esperado))
        return
    if not isinstance(got, (int, float)) or abs(float(got) - esperado) > tol:
        err.append('%s: %r ≠ esperado %r' % (etiqueta, got, esperado))


def _escenario(destino, cambios, err, nombre):
    """Hoja principal con valores cambiados vía pycel (compilador nuevo: sin estado de otros escenarios)."""
    from pycel import ExcelCompiler
    xl = ExcelCompiler(filename=destino)
    e = OrderedDict(_DEF)
    for k, v in cambios.items():
        xl.set_value("'%s'!B%d" % (H1, K[k]), v)
        e[k] = v
    esperado = calcular(e)
    n = 0
    for k, v in esperado.items():
        n += 1
        _cmp(err, '%s · %s!B%d (%s) pycel' % (nombre, H1, S[k], k), xl.evaluate("'%s'!B%d" % (H1, S[k])), v)
    return esperado, xl, n


def autotest(destino):
    from pycel import ExcelCompiler
    err, n = [], 0
    xl = ExcelCompiler(filename=destino)
    wv = openpyxl.load_workbook(destino, data_only=True)
    wf = openpyxl.load_workbook(destino)
    defecto = OrderedDict(_DEF)
    # 1) hoja principal y los 3 ejemplos (pycel + caché)
    casos = [(H4, 'BCD'[j], e, nombre) for j, (nombre, e) in enumerate(EJEMPLOS.items())]
    casos.append((H1, 'B', defecto, 'hoja principal (valores por defecto)'))
    for hoja, col, e, nombre in casos:
        esperado = calcular(e)
        for k, v in esperado.items():
            c = '%s%d' % (col, S[k])
            n += 1
            _cmp(err, '%s · %s!%s (%s) pycel' % (nombre, hoja, c, k), xl.evaluate("'%s'!%s" % (hoja, c)), v)
            _cmp(err, '%s · %s!%s (%s) caché' % (nombre, hoja, c, k), wv[hoja][c].value, v)
        if esperado['margin_calc'] + 1e-9 < e['margin']:
            err.append('%s: margen calculado %.4f < objetivo %.4f' % (nombre, esperado['margin_calc'], e['margin']))
        if esperado['mix_check'] != 'OK':
            err.append('%s: la mezcla de carnes no suma 100 %%' % nombre)
        if any(e['price%d' % i] is not None for i in _LINEAS):
            err.append('%s: precio de carne tecleado (D111: ninguno)' % nombre)
    # 2) los 3 ejemplos a ±10 % del precio de su línea (research §8) — precio por invitado y con las tarifas repartidas
    _wsx, f_equiv = hoja_ejemplos(openpyxl.Workbook())
    for j, (nombre, objetivo) in enumerate(OBJETIVO_EJ.items()):
        r = calcular(EJEMPLOS[nombre])
        equiv = r['subtotal'] / r['billable']
        for lab, v in (('precio por invitado', r['price']), ('con tarifas', equiv)):
            if abs(v / objetivo - 1) > 0.10 + 1e-9:
                err.append('%s: %s %.2f fuera de ±10 %% de %.2f' % (nombre, lab, v, objetivo))
        c = '%s%d' % ('BCD'[j], f_equiv)
        n += 1
        _cmp(err, 'equivalente %s pycel' % c, xl.evaluate("'%s'!%s" % (H4, c)), equiv)
        _cmp(err, 'equivalente %s caché' % c, wv[H4][c].value, equiv)
    # 3) Meat Yield & Smoke Schedule (pycel + caché)
    for i, (nombre, x) in enumerate(calcular_yield()):
        r = Y_F0 + i
        for k, v in x.items():
            col = Y_COL[k]
            n += 1
            tol = 1e-6
            _cmp(err, 'yield %s %s%d pycel' % (nombre, col, r), xl.evaluate("'%s'!%s%d" % (H2, col, r)), v, tol)
            _cmp(err, 'yield %s %s%d caché' % (nombre, col, r), wv[H2]['%s%d' % (col, r)].value, v, tol)
        if CORTES[i][3] < 100 or not wf[H2]['G%d' % r].value or '°F' not in wf[H2]['G%d' % r].value:
            err.append('yield %s: calidad / mínimo del Food Code sin rellenar' % nombre)
    # 4) Shopping List by Guest Count (pycel + caché), con los valores por defecto
    _ws3, filas = hoja_compra(openpyxl.Workbook())
    for j, col in enumerate(C_COLS):
        esperado = calcular_compra(defecto)[j]
        for k in C_CLAVES:
            c = '%s%d' % (col, filas[k])
            n += 1
            _cmp(err, 'compra %s (%s) pycel' % (c, k), xl.evaluate("'%s'!%s" % (H3, c)), esperado[k], 1e-6)
            _cmp(err, 'compra %s (%s) caché' % (c, k), wv[H3][c].value, esperado[k], 1e-6)
    # 5) escenarios en la hoja principal (solo pycel): paquete live-fire, interruptor del live pit, precios tecleados
    e_live, _x, m = _escenario(destino, {'package': LIVE}, err, 'live-fire')
    n += m
    if not e_live['price'] > calcular(defecto)['price'] or e_live['oz_guest'] != PAQUETES[LIVE][0]:
        err.append('paquete live-fire sin efecto')
    e_pit, _x, m = _escenario(destino, {'live_pit': 'Yes'}, err, 'live pit')
    n += m
    if abs(e_pit['subtotal'] - calcular(defecto)['subtotal'] - _DEF['live_pit_fee']) > 1e-6:
        err.append('interruptor del live pit sin efecto')
    precios = {'price1': 4.0, 'price2': 2.5, 'price3': 4.5, 'price4': 2.0, 'price5': 3.5}      # solo para la prueba
    e_carne, xlc, m = _escenario(destino, precios, err, 'precios de carne tecleados')
    n += m
    if e_carne['missing'] != 0 or not e_carne['c_meat'] > 0 or e_carne['meat_status'] == calcular(defecto)['meat_status']:
        err.append('precios de carne tecleados sin efecto')
    for k in ('raw1', 'pieces1', 'wood_lb'):                                     # las libras no dependen del precio
        if abs(e_carne[k] - calcular(defecto)[k]) > 1e-9:
            err.append('las libras cambian con el precio (%s)' % k)
    # 6) interruptor del sales tax sobre el service charge
    e = dict(EJEMPLOS['Wedding, 130 guests, full-service + live pit on site'])
    if not calcular(dict(e, tax_sc='Yes'))['tax_amt'] > calcular(dict(e, tax_sc='No'))['tax_amt']:
        err.append('interruptor Yes/No del sales tax sin efecto')
    # higiene del fichero (censo-entregables + D20 + D8)
    p = wf.properties
    if p.creator != 'AI Chef Pro' or not (p.title or '').startswith(TITULO) or M.PLANES['bbq']['url'] != p.description:
        err.append('docProps D20: %r / %r / %r' % (p.creator, p.title, p.description))
    textos = []
    for ws in wf.worksheets:
        if ws.page_setup.paperSize != 1 or not ws.sheet_properties.pageSetUpPr.fitToPage or \
                ws.oddFooter.center.text != M.FOOTER_EN:
            err.append('%s: impresión (Letter + ajuste + pie)' % ws.title)
        if not ws.protection.sheet:
            err.append('%s: hoja sin proteger' % ws.title)
        for row in ws.iter_rows():
            for c in row:
                v = c.value
                if v == '':
                    err.append('%s!%s: cadena vacía' % (ws.title, c.coordinate))
                if isinstance(v, str) and c.data_type != 'f':
                    textos.append((ws.title, c.coordinate, v))
                if isinstance(v, str) and c.data_type == 'f':
                    textos += [(ws.title, c.coordinate + ' (fórmula)', lit) for lit in re.findall(r'"((?:[^"]|"")*)"', v)
                               if any(ch.isalpha() for ch in lit)]
                if c.data_type == 'f' and wv[ws.title][c.coordinate].value is None:
                    err.append('%s!%s: fórmula sin caché' % (ws.title, c.coordinate))
                if c.fill.fgColor.rgb in ('00E8F5E9', 'FFE8F5E9') and c.protection.locked:
                    err.append('%s!%s: verde bloqueada' % (ws.title, c.coordinate))
                if isinstance(v, str) and c.data_type != 'f' and '€' in v:
                    err.append('%s!%s: «€»' % (ws.title, c.coordinate))
                fmt = c.number_format or ''
                if '€' in fmt or '$' in fmt:
                    err.append('%s!%s: símbolo de moneda en el formato (D8)' % (ws.title, c.coordinate))
        for dv in ws.data_validations.dataValidation:
            for t in (dv.error, dv.errorTitle, dv.prompt, dv.promptTitle):
                if t:
                    textos.append((ws.title, 'DV %s' % dv.sqref, t))
    for campo in ('title', 'subject', 'keywords', 'description'):
        textos.append(('docProps', campo, getattr(p, campo) or ''))
    for h, c, t in textos:
        r = [x for x in M.restos_espanol(t) if not x[0].startswith('oficio')]
        if r:
            err.append('%s!%s: restos de español %s' % (h, c, r[:1]))
        if M.no_latinos(t):
            err.append('%s!%s: caracteres no latinos' % (h, c))
        if M.cifras_prohibidas(t):
            err.append('%s!%s: cifra/término prohibido %s' % (h, c, M.cifras_prohibidas(t)[:1]))
        if re.search(r'(?i)\bvalue\b.*\$\d|\bsave \$|total value|\bwas \$|\d+\s?% off\b', t):
            err.append('%s!%s: ancla de precio' % (h, c))
    # temperaturas: la de calidad y el mínimo seguro nunca en la misma columna (SPEC §0)
    cab = [wf[H2].cell(Y_F0 - 1, j).value for j in range(1, 16)]
    if not any('Quality' in (x or '') for x in cab) or not any('Food Code' in (x or '') for x in cab):
        err.append('%s: falta la columna de calidad o la del mínimo del Food Code' % H2)
    # ningún precio de carne tecleado (D111): las celdas de $/lb, en blanco
    for i in _LINEAS:
        for hoja, cols in ((H1, 'B'), (H4, 'BCD')):
            for col in cols:
                if wf[hoja]['%s%d' % (col, K['price%d' % i])].value is not None:
                    err.append('%s!%s%d: precio de carne tecleado' % (hoja, col, K['price%d' % i]))
    hojas = [ws.title for ws in wf.worksheets]
    if hojas != [H1, H2, H3, H4, H5]:
        err.append('pestañas %r' % hojas)
    for x in err[:80]:
        print('ERROR', x)
    print('autotest calculadora: %d comprobaciones numéricas (pycel + caché) en %s + live-fire, live pit y precios de '
          'carne en la hoja principal · %d textos · %s' % (n, ', '.join(HOJAS_CALCULO), len(textos),
                                                           'VERDE' if not err else 'ROJO (%d errores)' % len(err)))
    for nombre, e in list(EJEMPLOS.items()) + [('Main tab (defaults)', defecto), ('Main tab, live-fire', dict(defecto,
                                                                                                          package=LIVE)),
                                               ('Main tab, meat prices typed (test only)', dict(defecto, **precios))]:
        r = calcular(e)
        print('  %-52s %3d inv · %5.1f lb crudas · coste %8.2f · precio/inv %6.2f (/lb coc. %6.2f) · subtotal %9.2f · '
              'total %9.2f · margen %.1f%% · %s' % (nombre, e['guests'], r['raw_lb'], r['c_total'], r['price'],
                                                     r['price_lb'], r['subtotal'], r['total'], r['margin_calc'] * 100,
                                                     r['meat_status'][:8]))
    c = calcular_compra(defecto)
    print('  compra: ' + ' · '.join('%d inv: %.0f lb crudas, %d packers, %d butts, %d hot boxes' % (
        x['guests'], x['raw_tot'], x['pz1'], x['pz2'], x['boxes']) for x in c))
    return 1 if err else 0


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--autotest', action='store_true', help='solo el autotest sobre el fichero existente')
    ap.add_argument('--salida', default=DESTINO)
    a = ap.parse_args()
    if not a.autotest:
        if generar(a.salida):
            print('ERROR: fallos de pycel en inject_cache')
            return 1
    return autotest(a.salida)


if __name__ == '__main__':
    sys.exit(main())
