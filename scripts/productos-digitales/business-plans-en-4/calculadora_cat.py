#!/usr/bin/env python3
"""
calculadora_cat.py — Catering Pricing Calculator (Cost per Guest to Quote), D76 del Catering Business Plan Kit (EN).

Código NUEVO (no pasa por el circuito ES→EN: F1-inventario-delta §4): coste → precio por invitado. Entradas verdes con
validación (listas para el estilo de servicio y el interruptor del sales tax sobre el service charge), salidas
calculadas (el margen bruto se CALCULA, no se teclea), hoja «3 Example Quotes» calculada por las mismas fórmulas,
hoja «Instructions», docProps D20, US Letter, sin anclas de precio. Todo valor por defecto es «Kit estimate».
Corre en el VPS (openpyxl + pycel). Sesión Claude Code, 11-oct-2026 (F2, paso 1).

    python3 bp4.py calculadora              # genera + inject_cache + autotest
    python3 bp4.py calculadora --autotest   # solo el autotest sobre el fichero ya generado
"""
import argparse
import math
import os
import re
import sys
from collections import OrderedDict

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill, Protection
from openpyxl.worksheet.datavalidation import DataValidation

AQUI = os.path.dirname(os.path.abspath(__file__))
PD = os.path.dirname(AQUI)
REPO = os.path.abspath(os.path.join(AQUI, '..', '..', '..'))
sys.path.insert(0, os.path.join(PD, 'business-plans-en'))
sys.path.insert(0, PD)
os.environ.setdefault('BP_CONJUNTO', 'cat')
import mapas as M                                                                  # noqa: E402

SLUG = 'catering-business-plan'
FICHERO = 'catering-pricing-calculator.xlsx'
TITULO = 'Catering Pricing Calculator (Cost per Guest to Quote)'
DESTINO = os.path.join(REPO, 'astro-site', 'public', 'dl', SLUG, FICHERO)
H1, H2, H3 = 'Pricing Calculator', '3 Example Quotes', 'Instructions'
VERDE = PatternFill('solid', fgColor='E8F5E9')
GRIS = PatternFill('solid', fgColor='F2F2F2')
F = 'Calibri'
ESTILOS = ['Drop-off', 'Buffet', 'Family-style', 'Plated', 'Stations']
MONEDA, ENTERO, PCT, PCT1 = '#,##0.00', '#,##0', '0%', '0.0%'

# --------------------------------------------------------------------------
# Entradas: (clave, rótulo, formato, valor por defecto, nota). Fila = 5 + índice
# --------------------------------------------------------------------------
ENTRADAS = [
    ('guests', 'Guests expected', ENTERO, 60, 'Whole number. Food, beverage and disposables are costed on these guests.'),
    ('minimum', 'Guest minimum in your contract', ENTERO, 25,
     'If fewer guests confirm, you bill the minimum (billable guests below). 0 = no minimum.'),
    ('style', 'Service style', '@', 'Buffet',
     'Drop-off, Buffet, Family-style, Plated or Stations. It sets the suggested event staff (table at the bottom).'),
    ('food', 'Food cost per guest (from your recipe costing)', MONEDA, 18.0,
     'Ingredient cost of the menu per guest. Kit estimate: use your own recipe costing.'),
    ('beverage', 'Non-alcoholic beverage package per guest', MONEDA, 2.5,
     'Kit estimate. Alcohol is not included: the client buys it and certified bartenders serve it.'),
    ('disposables', 'Disposables per guest (plates, cutlery, packaging)', MONEDA, 0.5, 'Kit estimate.'),
    ('staff', 'Event staff — people', ENTERO, 5, 'Servers, bartenders, attendants and the chef on site.'),
    ('hours', 'Event staff — hours each (travel, setup, service, pack-down)', '#,##0.0', 7, 'Kit estimate.'),
    ('wage', 'Event staff — average hourly wage', MONEDA, 20.0,
     'Kit estimate (servers and bartenders usually earn $18-$25 an hour: restroworks.com).'),
    ('payroll', "Employer payroll taxes and workers' comp (% of wages)", PCT, 0.10,
     'Kit estimate. Event staff are W-2 employees, not 1099 contractors.'),
    ('rentals', 'Rentals passed through (tables, linens, china)', MONEDA, 0.0,
     'What the rental company charges you for this event.'),
    ('miles', 'Round-trip miles', '#,##0', 30, 'Commissary → venue → commissary.'),
    ('per_mile', 'Cost per mile', MONEDA, 0.70,
     'Kit estimate: use your real cost per mile or the current IRS standard mileage rate.'),
    ('margin', 'Target gross margin on the food & service subtotal', PCT, 0.50,
     'Share of the subtotal left after the direct costs above. Kit estimate.'),
    ('sc', 'Service charge (% of subtotal)', PCT, 0.20,
     'A mandatory service charge is wages, not a tip (IRS Rev. Rul. 2012-18). Kit estimate.'),
    ('tax', 'Sales tax rate (state + local)', PCT1, 0.08, 'Kit estimate: use your own state + local rate.'),
    ('tax_sc', 'Does your state tax the service charge?', '@', 'Yes',
     'CA: a mandatory service charge is taxable. NY: not taxable only if it is separately stated as a gratuity and '
     'all of it goes to the staff. TX: not taxable up to 20% if separately stated and distributed to qualified '
     'employees. Check your state.'),
    ('gratuity', 'Suggested voluntary gratuity (% of subtotal; NOT in the price)', PCT, 0.0,
     'Optional. A voluntary tip is the client\'s decision and goes to the staff: it stays out of the total.'),
    ('deposit', 'Deposit to book the date (% of total)', PCT, 0.35,
     'Kit estimate (25-50% at signing is common); the balance is due before the event.'),
]
K = {k: 5 + i for i, (k, *_r) in enumerate(ENTRADAS)}                  # clave → fila de la hoja 1
F_CALC = 5 + len(ENTRADAS) + 1                                         # cabecera del bloque de cálculo
SALIDAS = [
    ('billable', 'Billable guests', ENTERO, 'max({guests},{minimum})', 'The larger of guests expected and the minimum.'),
    ('c_food', 'Food, beverage and disposables', MONEDA, '({food}+{beverage}+{disposables})*{guests}', ''),
    ('c_staff', 'Event staff cost (wages + payroll taxes)', MONEDA, '{staff}*{hours}*{wage}*(1+{payroll})', ''),
    ('c_other', 'Rentals and mileage', MONEDA, '{rentals}+{miles}*{per_mile}', ''),
    ('c_total', 'Total direct cost', MONEDA, '{c_food}+{c_staff}+{c_other}', ''),
    ('c_guest', 'Direct cost per billable guest', MONEDA, '{c_total}/{billable}', ''),
    ('price', 'PRICE PER GUEST (rounded up to the next 0.50)', MONEDA, 'roundup({c_total}/(1-{margin})/{billable}*2)/2',
     'Excl. service charge and sales tax.'),
    ('subtotal', 'Food & service subtotal', MONEDA, '{price}*{billable}', 'Price per guest × billable guests.'),
    ('sc_amt', 'Service charge', MONEDA, '{subtotal}*{sc}', ''),
    ('tax_amt', 'Sales tax', MONEDA, '({subtotal}+if({tax_sc}=="Yes",{sc_amt},0))*{tax}',
     'On the subtotal, plus the service charge if your state taxes it.'),
    ('total', 'TOTAL TO QUOTE', MONEDA, '{subtotal}+{sc_amt}+{tax_amt}', ''),
    ('dep_amt', 'Deposit to book the date', MONEDA, '{total}*{deposit}', ''),
    ('balance', 'Balance due before the event', MONEDA, '{total}-{dep_amt}', ''),
    ('profit', 'Gross profit on the subtotal', MONEDA, '{subtotal}-{c_total}', ''),
    ('margin_calc', 'Gross margin (calculated)', PCT1, 'if({subtotal}==0,0,{profit}/{subtotal})',
     'Calculated from the costs above, never typed. It is at or above the target because the price is rounded up.'),
    ('grat_amt', 'Suggested voluntary gratuity (not in the total)', MONEDA, '{subtotal}*{gratuity}', ''),
    ('staff_sug', 'Suggested event staff for this style (Kit estimate)', ENTERO, 'sugerido', 'From the table below.'),
]
S = {k: F_CALC + 1 + i for i, (k, *_r) in enumerate(SALIDAS)}
F_TABLA = F_CALC + 1 + len(SALIDAS) + 1                                # cabecera de la tabla de estilos
RATIO = OrderedDict([('Drop-off', 0), ('Buffet', 25), ('Family-style', 15), ('Plated', 12), ('Stations', 20)])

EJEMPLOS = OrderedDict([
    ('Corporate drop-off lunch', dict(guests=40, minimum=20, style='Drop-off', food=9.0, beverage=1.5,
                                      disposables=1.25, staff=1, hours=3, wage=18.0, payroll=0.10, rentals=0.0,
                                      miles=20, per_mile=0.70, margin=0.55, sc=0.0, tax=0.08, tax_sc='Yes',
                                      gratuity=0.0, deposit=0.0)),
    ('Wedding buffet', dict(guests=120, minimum=75, style='Buffet', food=24.0, beverage=3.0, disposables=0.5,
                            staff=8, hours=7, wage=20.0, payroll=0.10, rentals=1800.0, miles=60, per_mile=0.70,
                            margin=0.50, sc=0.20, tax=0.08, tax_sc='Yes', gratuity=0.0, deposit=0.35)),
    ('Plated dinner', dict(guests=60, minimum=40, style='Plated', food=16.0, beverage=3.0, disposables=0.0, staff=6,
                           hours=6, wage=20.0, payroll=0.10, rentals=600.0, miles=30, per_mile=0.70, margin=0.50,
                           sc=0.20, tax=0.08, tax_sc='No', gratuity=0.0, deposit=0.35)),
])
EJ_NOTAS = OrderedDict([
    ('Corporate drop-off lunch', 'Drop-off: one driver delivers and sets up; no service charge; the company pays the '
                                 'invoice, often within 30 days (no deposit).'),
    ('Wedding buffet', 'Buffet with servers and a bartender for the client\'s alcohol; rentals passed through; 35% '
                       'deposit to book the date.'),
    ('Plated dinner', 'Plated service needs more servers per guest; this example assumes a state that does not tax '
                      'the service charge (switch set to No).'),
])


# --------------------------------------------------------------------------
# Lógica en Python (la misma que las fórmulas): la usa el autotest como valor esperado
# --------------------------------------------------------------------------
def calcular(e):
    r = OrderedDict()
    r['billable'] = max(e['guests'], e['minimum'])
    r['c_food'] = (e['food'] + e['beverage'] + e['disposables']) * e['guests']
    r['c_staff'] = e['staff'] * e['hours'] * e['wage'] * (1 + e['payroll'])
    r['c_other'] = e['rentals'] + e['miles'] * e['per_mile']
    r['c_total'] = r['c_food'] + r['c_staff'] + r['c_other']
    r['c_guest'] = r['c_total'] / r['billable']
    r['price'] = math.ceil(round(r['c_total'] / (1 - e['margin']) / r['billable'] * 2, 9)) / 2
    r['subtotal'] = r['price'] * r['billable']
    r['sc_amt'] = r['subtotal'] * e['sc']
    r['tax_amt'] = (r['subtotal'] + (r['sc_amt'] if e['tax_sc'] == 'Yes' else 0)) * e['tax']
    r['total'] = r['subtotal'] + r['sc_amt'] + r['tax_amt']
    r['dep_amt'] = r['total'] * e['deposit']
    r['balance'] = r['total'] - r['dep_amt']
    r['profit'] = r['subtotal'] - r['c_total']
    r['margin_calc'] = 0 if r['subtotal'] == 0 else r['profit'] / r['subtotal']
    r['grat_amt'] = r['subtotal'] * e['gratuity']
    ratio = RATIO[e['style']]
    r['staff_sug'] = 1 if ratio == 0 else math.ceil(e['guests'] / ratio)
    return r


def a_excel(expr, ref):
    """Plantilla Python-ish → fórmula Excel con las celdas de `ref` (clave → coordenada)."""
    f = expr.replace('max(', 'MAX(').replace('roundup(', 'ROUNDUP(').replace('if(', 'IF(').replace('==', '=')
    f = re.sub(r'\{(\w+)\}', lambda m: ref[m.group(1)], f)
    return f.replace('*2)/2', '*2,0)/2')


def formula_sugerido(ref_estilo, ref_guests):
    rng_e = "'%s'!$A$%d:$A$%d" % (H1, F_TABLA + 1, F_TABLA + len(RATIO))
    rng_r = "'%s'!$B$%d:$B$%d" % (H1, F_TABLA + 1, F_TABLA + len(RATIO))
    rat = 'INDEX(%s,MATCH(%s,%s,0))' % (rng_r, ref_estilo, rng_e)
    return '=IFERROR(IF(%s=0,1,ROUNDUP(%s/%s,0)),"—")' % (rat, ref_guests, rat)


# --------------------------------------------------------------------------
# Construcción
# --------------------------------------------------------------------------
def _f(ws, coord, v, b=False, sz=11, i=False, fmt=None, wrap=False, fill=None, locked=True, color=None):
    c = ws[coord]
    c.value = v
    c.font = Font(name=F, size=sz, bold=b, italic=i, color=color)
    if fmt:
        c.number_format = fmt
    c.alignment = Alignment(wrap_text=wrap, vertical='top')
    if fill:
        c.fill = fill
    c.protection = Protection(locked=locked)
    return c


def _imprimir(ws, horizontal=False):
    ws.page_setup.paperSize = M.PAPEL_EN                                         # 1 = US Letter
    ws.page_setup.orientation = 'landscape' if horizontal else 'portrait'
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.oddFooter.center.text = M.FOOTER_EN
    ws.print_options.horizontalCentered = True
    ws.protection.sheet = True                                                   # sin contraseña (D20)
    ws.sheet_view.showGridLines = False


def _dv(ws, tipo, celdas, f1=None, f2=None, op=None, titulo='Invalid value', msg=''):
    dv = DataValidation(type=tipo, formula1=f1, formula2=f2, operator=op, allow_blank=False, showErrorMessage=True,
                        errorTitle=titulo, error=msg)
    ws.add_data_validation(dv)
    for c in celdas:
        dv.add(c)


def hoja_calculadora(wb):
    ws = wb.active
    ws.title = H1
    for col, w in (('A', 58), ('B', 16), ('C', 78)):
        ws.column_dimensions[col].width = w
    _f(ws, 'A1', 'CATERING PRICING CALCULATOR — COST PER GUEST TO QUOTE', b=True, sz=14)
    _f(ws, 'A2', 'Type in the GREEN cells; everything else is calculated. Amounts are in your currency and exclude '
                 'sales tax unless the row says otherwise. Every default value is a kit estimate: replace it with '
                 'your own numbers.', wrap=True)
    ws.merge_cells('A2:C2')
    ws.row_dimensions[2].height = 32
    for col, t in zip('ABC', ('INPUTS', 'Value', 'Notes')):
        _f(ws, '%s4' % col, t, b=True, fill=GRIS)
    for k, rot, fmt, val, nota in ENTRADAS:
        f = K[k]
        _f(ws, 'A%d' % f, rot)
        _f(ws, 'B%d' % f, val, fmt=fmt if fmt != '@' else None, fill=VERDE, locked=False)
        _f(ws, 'C%d' % f, nota, wrap=True)
    for col, t in zip('ABC', ('CALCULATION', 'Value', 'Notes')):
        _f(ws, '%s%d' % (col, F_CALC), t, b=True, fill=GRIS)
    ref = {k: '$B$%d' % f for k, f in K.items()}
    ref.update({k: '$B$%d' % f for k, f in S.items()})
    for k, rot, fmt, expr, nota in SALIDAS:
        f = S[k]
        fx = formula_sugerido(ref['style'], ref['guests']) if expr == 'sugerido' else '=' + a_excel(expr, ref)
        neg = k in ('price', 'total', 'margin_calc')
        _f(ws, 'A%d' % f, rot, b=neg)
        _f(ws, 'B%d' % f, fx, fmt=fmt, b=neg)
        if nota:
            _f(ws, 'C%d' % f, nota, wrap=True)
    for col, t in zip('ABC', ('GUESTS PER SERVER BY SERVICE STYLE (Kit estimate — edit)', 'Guests', 'Notes')):
        _f(ws, '%s%d' % (col, F_TABLA), t, b=True, fill=GRIS)
    notas = {'Drop-off': '0 = no servers: one person delivers and sets up.',
             'Plated': 'About 1 server per 10-12 guests.', 'Buffet': 'About 1 server per 20-25 guests.'}
    for i, (est, n) in enumerate(RATIO.items(), 1):
        _f(ws, 'A%d' % (F_TABLA + i), est)
        _f(ws, 'B%d' % (F_TABLA + i), n, fmt=ENTERO, fill=VERDE, locked=False)
        if est in notas:
            _f(ws, 'C%d' % (F_TABLA + i), notas[est], wrap=True)
    # validaciones (D76)
    _dv(ws, 'list', [ref['style'].replace('$', '')], f1='"%s"' % ','.join(ESTILOS), titulo='Invalid style',
        msg='Choose a service style from the list.')
    _dv(ws, 'list', [ref['tax_sc'].replace('$', '')], f1='"Yes,No"', titulo='Invalid entry', msg='Choose Yes or No.')
    _dv(ws, 'whole', ['B%d' % K['guests']], f1='1', f2='100000', op='between', titulo='Invalid count',
        msg='Enter a whole number of 1 or more.')
    _dv(ws, 'whole', ['B%d' % K[k] for k in ('minimum', 'staff')] +
        ['B%d' % (F_TABLA + i) for i in range(1, len(RATIO) + 1)], f1='0', f2='100000', op='between',
        titulo='Invalid count', msg='Enter a whole number of 0 or more.')
    _dv(ws, 'decimal', ['B%d' % K[k] for k in ('food', 'beverage', 'disposables', 'hours', 'wage', 'rentals', 'miles',
                                              'per_mile')], f1='0', op='greaterThanOrEqual', titulo='Invalid amount',
        msg='Enter a number greater than or equal to 0.')
    _dv(ws, 'decimal', ['B%d' % K[k] for k in ('payroll', 'sc', 'tax', 'gratuity', 'deposit')], f1='0', f2='1',
        op='between', titulo='Invalid percentage', msg='Enter a percentage between 0 and 1 (0.20 = 20%).')
    _dv(ws, 'decimal', ['B%d' % K['margin']], f1='0', f2='0.95', op='between', titulo='Invalid percentage',
        msg='Enter a target margin between 0 and 95% (0.50 = 50%).')
    ws.freeze_panes = 'A5'
    _imprimir(ws)
    return ws


def hoja_ejemplos(wb):
    ws = wb.create_sheet(H2)
    ws.column_dimensions['A'].width = 58
    for col in 'BCD':
        ws.column_dimensions[col].width = 24
    _f(ws, 'A1', '3 EXAMPLE QUOTES — calculated with the same formulas as the "%s" tab' % H1, b=True, sz=14)
    _f(ws, 'A2', 'Kit estimates for illustration only, not market prices: change the inputs in "%s" to price your own '
                 'events.' % H1, wrap=True)
    ws.merge_cells('A2:D2')
    ws.row_dimensions[2].height = 32
    _f(ws, 'A4', 'INPUTS', b=True, fill=GRIS)
    for j, nombre in enumerate(EJEMPLOS):
        col = 'BCD'[j]
        _f(ws, '%s4' % col, nombre, b=True, fill=GRIS, wrap=True)
    for k, rot, fmt, val, nota in ENTRADAS:
        _f(ws, 'A%d' % K[k], rot)
        for j, e in enumerate(EJEMPLOS.values()):
            _f(ws, '%s%d' % ('BCD'[j], K[k]), e[k], fmt=fmt if fmt != '@' else None)
    _f(ws, 'A%d' % F_CALC, 'RESULT', b=True, fill=GRIS)
    for j in range(3):
        _f(ws, '%s%d' % ('BCD'[j], F_CALC), list(EJEMPLOS)[j], b=True, fill=GRIS, wrap=True)
    for k, rot, fmt, expr, nota in SALIDAS:
        neg = k in ('price', 'total', 'margin_calc')
        _f(ws, 'A%d' % S[k], rot, b=neg)
        for j in range(3):
            col = 'BCD'[j]
            ref = {kk: '$%s$%d' % (col, f) for kk, f in K.items()}
            ref.update({kk: '$%s$%d' % (col, f) for kk, f in S.items()})
            fx = formula_sugerido(ref['style'], ref['guests']) if expr == 'sugerido' else '=' + a_excel(expr, ref)
            _f(ws, '%s%d' % (col, S[k]), fx, fmt=fmt, b=neg)
    f = S[SALIDAS[-1][0]] + 2
    _f(ws, 'A%d' % f, 'WHAT EACH EXAMPLE SHOWS', b=True, fill=GRIS)
    for j, (nombre, nota) in enumerate(EJ_NOTAS.items()):
        _f(ws, 'A%d' % (f + 1 + j), nombre, b=True)
        _f(ws, 'B%d' % (f + 1 + j), nota, wrap=True)
        ws.merge_cells('B%d:D%d' % (f + 1 + j, f + 1 + j))
        ws.row_dimensions[f + 1 + j].height = 30
    ws.freeze_panes = 'B5'
    _imprimir(ws, horizontal=True)
    return ws


INSTRUCCIONES = [
    ('HOW TO USE THE CATERING PRICING CALCULATOR', True),
    ('This calculator turns your costs into a price per guest and a full quote: subtotal, service charge, sales tax, '
     'total, deposit and balance. You type only in the GREEN cells; the gross margin is calculated, never typed.', False),
    ('1. Enter the guests expected and the guest minimum in your contract: below the minimum you bill the minimum.', False),
    ('2. Choose the service style. It sets the suggested event staff from the table at the bottom of the first tab '
     '(guests per server by style: edit it to match your own standards).', False),
    ('3. Enter your food cost per guest from your recipe costing, the non-alcoholic beverage package and the '
     'disposables. Alcohol is not included: by default the client buys it and certified bartenders serve it.', False),
    ('4. Enter the event staff (people, hours each, average hourly wage, payroll taxes) and the rentals and mileage of '
     'the event.', False),
    ('5. Set your target gross margin on the food & service subtotal. The price per guest is the direct cost divided '
     'by (1 - target margin), rounded up to the next 0.50, so the calculated margin is at or above your target.', False),
    ('6. Set the service charge, your sales tax rate and the switch "Does your state tax the service charge?".', False),
    ('7. The deposit books the date; the balance is due before the event. The "3 Example Quotes" tab shows three '
     'complete quotes calculated with the same formulas.', False),
    ('SERVICE CHARGE, GRATUITY AND SALES TAX (US)', True),
    ('A mandatory service charge is not a tip: the IRS treats it as wages paid to the staff, with payroll taxes (Rev. '
     'Rul. 2012-18). A voluntary gratuity is the client\'s decision: this calculator keeps it out of the price and the '
     'total.', False),
    ('Sales tax on catering: in many states it applies to the whole catering invoice (food, staff, rentals). Whether '
     'it also applies to the service charge depends on the state: in California a mandatory service charge is '
     'taxable; in New York it is not taxable only when it is separately stated as a gratuity and all of it goes to '
     'the staff; in Texas it is not taxable up to 20% if separately stated and distributed to qualified employees. '
     'Check your state\'s rules and set the switch.', False),
    ('UK: catering for events and functions is standard-rated for VAT (20%, HMRC VAT Notice 709/1): enter 0.20 as the '
     'tax rate and include it in your quotes.', False),
    ('KIT ESTIMATES', True),
    ('Every default value and the three example quotes are kit estimates for illustration, not market prices. Replace '
     'them with your own recipe costing, wages, rentals and tax rate before you quote.', False),
]


def hoja_instrucciones(wb):
    ws = wb.create_sheet(H3)
    ws.column_dimensions['A'].width = 110
    f = 1
    for t, b in INSTRUCCIONES:
        _f(ws, 'A%d' % f, t, b=b, sz=13 if f == 1 else 11, wrap=True)
        if not b:
            ws.row_dimensions[f].height = max(16, 15 * math.ceil(len(t) / 120.0))
        f += 2 if b and f > 1 else 1
    f += 1
    for t in ('To edit a cell that is not green: Review → Unprotect Sheet (there is no password).',
              'More templates and kits in the catalog at ' + M.URL_TIENDA, M.FIRMA_EN, M.version_en('cat'),
              M.AVISO_EN):
        _f(ws, 'A%d' % f, t, wrap=True, i=t == M.AVISO_EN)
        f += 1
    _imprimir(ws)
    return ws


def docprops(wb):
    p = M.PLANES['cat']
    pr = wb.properties
    pr.title = '%s · %s' % (TITULO, p['producto'])
    pr.subject = '%s · v%s' % (p['producto'], M.VERSION)
    pr.keywords = p['keywords']
    pr.description = p['url']
    pr.category = 'AI Chef Pro · Digital products'
    pr.creator = 'AI Chef Pro'
    pr.lastModifiedBy = 'AI Chef Pro'


def generar(destino):
    wb = openpyxl.Workbook()
    hoja_calculadora(wb)
    hoja_ejemplos(wb)
    hoja_instrucciones(wb)
    docprops(wb)
    os.makedirs(os.path.dirname(destino), exist_ok=True)
    wb.save(destino)
    sys.path.insert(0, PD)
    import inject_cache as IC                                                    # noqa: E402
    nf, inj, fail = IC.inject(destino)
    print('%s: formulas=%d cache_inyectado=%d fallos_pycel=%d' % (os.path.relpath(destino, REPO), nf, inj, fail))
    return fail


# --------------------------------------------------------------------------
# Autotest (D76): pycel recalcula los 3 ejemplos y la hoja principal = lo esperado; caché = lo esperado; 0 errores
# --------------------------------------------------------------------------
def autotest(destino):
    from pycel import ExcelCompiler
    err, n = [], 0
    xl = ExcelCompiler(filename=destino)
    wv = openpyxl.load_workbook(destino, data_only=True)
    wf = openpyxl.load_workbook(destino)
    casos = [(H2, 'BCD'[j], e, nombre) for j, (nombre, e) in enumerate(EJEMPLOS.items())]
    casos.append((H1, 'B', OrderedDict((k, v) for k, _r, _f2, v, _n in ENTRADAS), 'hoja principal (valores por defecto)'))
    for hoja, col, e, nombre in casos:
        esperado = calcular(e)
        for k, v in esperado.items():
            c = '%s%d' % (col, S[k])
            got = xl.evaluate("'%s'!%s" % (hoja, c))
            cache = wv[hoja][c].value
            n += 1
            for lab, x in (('pycel', got), ('caché', cache)):
                if not isinstance(x, (int, float)) or abs(float(x) - v) > 0.005:
                    err.append('%s · %s %s (%s): %s %r ≠ esperado %r' % (nombre, hoja, c, k, lab, x, v))
        if esperado['margin_calc'] + 1e-9 < e['margin']:
            err.append('%s: margen calculado %.4f < objetivo %.4f' % (nombre, esperado['margin_calc'], e['margin']))
        if 'Yes' == e['tax_sc'] and abs(esperado['tax_amt'] - (esperado['subtotal'] + esperado['sc_amt']) * e['tax']) > 0.005:
            err.append('%s: el interruptor del sales tax no grava el service charge' % nombre)
    # el interruptor cambia el resultado (misma entrada con Yes / No)
    e = dict(EJEMPLOS['Wedding buffet'])
    si, no = calcular(dict(e, tax_sc='Yes')), calcular(dict(e, tax_sc='No'))
    if not si['tax_amt'] > no['tax_amt']:
        err.append('interruptor Yes/No sin efecto')
    # higiene del fichero (censo-entregables + D20 + D8)
    p = wf.properties
    if p.creator != 'AI Chef Pro' or not (p.title or '').startswith(TITULO) or M.PLANES['cat']['url'] != p.description:
        err.append('docProps D20: %r / %r / %r' % (p.creator, p.title, p.description))
    textos = []
    for ws in wf.worksheets:
        if ws.page_setup.paperSize != 1 or not ws.sheet_properties.pageSetUpPr.fitToPage or \
                ws.oddFooter.center.text != M.FOOTER_EN:
            err.append('%s: impresión (Letter + ajuste + pie)' % ws.title)
        if not ws.protection.sheet:
            err.append('%s: hoja sin proteger' % ws.title)
        for row in ws.iter_rows():
            for c in row:
                v = c.value
                if v == '':
                    err.append('%s!%s: cadena vacía' % (ws.title, c.coordinate))
                if isinstance(v, str) and c.data_type != 'f':
                    textos.append((ws.title, c.coordinate, v))
                if c.data_type == 'f' and wv[ws.title][c.coordinate].value is None:
                    err.append('%s!%s: fórmula sin caché' % (ws.title, c.coordinate))
                if c.fill.fgColor.rgb in ('00E8F5E9', 'FFE8F5E9') and c.protection.locked:
                    err.append('%s!%s: verde bloqueada' % (ws.title, c.coordinate))
                if isinstance(v, str) and c.data_type != 'f' and '€' in v:
                    err.append('%s!%s: «€»' % (ws.title, c.coordinate))
                fmt = c.number_format or ''
                if '€' in fmt or '$' in fmt:
                    err.append('%s!%s: símbolo de moneda en el formato (D8)' % (ws.title, c.coordinate))
    for h, c, t in textos:
        r = [x for x in M.restos_espanol(t) if not x[0].startswith('oficio')]
        if r:
            err.append('%s!%s: restos de español %s' % (h, c, r[:1]))
        if M.no_latinos(t):
            err.append('%s!%s: caracteres no latinos' % (h, c))
        if M.cifras_prohibidas(t):
            err.append('%s!%s: cifra/término prohibido %s' % (h, c, M.cifras_prohibidas(t)[:1]))
        if re.search(r'(?i)\bvalue\b.*\$\d|\bsave \$|total value|\bwas \$', t):
            err.append('%s!%s: ancla de precio' % (h, c))
    hojas = [ws.title for ws in wf.worksheets]
    if hojas != [H1, H2, H3]:
        err.append('pestañas %r' % hojas)
    for x in err:
        print('ERROR', x)
    print('autotest calculadora: %d comprobaciones numéricas (pycel + caché) sobre 3 ejemplos + la hoja principal · '
          '%d textos · %s' % (n * 2, len(textos), 'VERDE' if not err else 'ROJO (%d errores)' % len(err)))
    for nombre, e in EJEMPLOS.items():
        r = calcular(e)
        print('  %-26s %3d inv · precio %7.2f · subtotal %9.2f · service charge %8.2f · tax %7.2f · total %9.2f · '
              'margen %.1f%% · personal sugerido %d' % (nombre, r['billable'], r['price'], r['subtotal'], r['sc_amt'],
                                                       r['tax_amt'], r['total'], r['margin_calc'] * 100, r['staff_sug']))
    return 1 if err else 0


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--autotest', action='store_true', help='solo el autotest sobre el fichero existente')
    ap.add_argument('--salida', default=DESTINO)
    a = ap.parse_args()
    if not a.autotest:
        if generar(a.salida):
            print('ERROR: fallos de pycel en inject_cache')
            return 1
    return autotest(a.salida)


if __name__ == '__main__':
    sys.exit(main())
