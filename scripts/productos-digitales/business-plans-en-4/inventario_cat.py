#!/usr/bin/env python3
"""
inventario_cat.py — Catering Business Plan Kit (EN) · F1 (sesión Claude Code, 11-oct-2026).

Inventario MEDIDO de la fuente ES de LÍNEA B `plan-catering-tematico-eventos` (solo los 4 ficheros en alcance: plan
docx, plan financiero xlsx, calculadora de precios xlsx, checklist xlsx) y comparación con los 5 libros EN de línea A
ya publicados (FT, CAF, BAR, REST, PAN) para elegir el gemelo estructural. SOLO LECTURA: no escribe nada salvo el JSON.

Qué mide:
  · sha256 de los 11 ficheros ES (4 en alcance + 7 fuera);
  · por xlsx ES: hojas, rangos, fórmulas, fórmulas con CONSTANTES numéricas dentro (tasas tecleadas en la fórmula),
    textos (celdas y palabras), DV, CF, merges, protección, rellenos de entrada; valores clave de la caché;
  · coherencias del caso ES (tickets ponderados de «Mix B2C-B2B» frente a «Proyección 3 años», márgenes, PE,
    inversión del xlsx frente al docx y la ficha);
  · checklist ES: tareas por fase, clasificación (ES_SOLO / ADAPTAR / MULTICOCINA / directo) y reparto propuesto a las
    6 fases del molde EN;
  · docx ES: párrafos, estilos, secciones (por texto «N. »), palabras por sección, tablas, importes en €;
  · libros EN publicados: hojas, fórmulas, rótulos de B4/B5/B6/B24 de «0. Assumptions», partidas tecleadas de
    «Startup Costs», puestos de «Staffing», fijos del P&L, fila de estacionalidad, tareas por fase del checklist y
    párrafos/encabezados/palabras del docx;
  · en el gemelo FT EN: qué fórmulas leen cada driver de actividad ($B$4 clientes/día, $B$5 ticket, $B$6 días),
    para saber cuántas cambian si se reinterpreta el motor, y cuántas celdas de texto mencionan la actividad del FT.

    /root/venv-guias/bin/python inventario_cat.py [--repo DIR] [--salida inventario_cat.json]

Corre en el VPS (openpyxl + python-docx), en serie, unos segundos.
"""
import argparse
import collections
import hashlib
import json
import os
import re

import docx
import openpyxl

AQUI = os.path.dirname(os.path.abspath(__file__))
DL = 'astro-site/public/dl'
ES_DIR = 'plan-catering-tematico-eventos'
ES_ALCANCE = collections.OrderedDict([
    ('CATP', 'plan-financiero-catering-tematico-eventos.xlsx'),
    ('CATPR', 'calculadora-pricing-catering-tematico-eventos.xlsx'),
    ('CATC', 'checklist-apertura-catering-tematico-eventos.xlsx'),
    ('CATD', 'plan-de-negocio-catering-tematico-eventos.docx'),
])
FICHA_ES = 'astro-site/src/data/productos/planes/plan-catering-tematico-eventos.ts'
EN = collections.OrderedDict([          # corto → (carpeta, prefijo de fichero)
    ('FT', ('food-truck-business-plan', 'food-truck')),
    ('CAF', ('coffee-shop-business-plan', 'coffee-shop')),
    ('BAR', ('bar-business-plan', 'bar')),
    ('REST', ('restaurant-business-plan', 'restaurant')),
    ('PAN', ('bakery-business-plan', 'bakery')),
])
DOCX_FT_ES = ('plan-negocio-food-truck', 'plan-de-negocio-food-truck.docx')
RX_NUM_EN_FX = re.compile(r'(?<![A-Z$\d.])(\d+\.\d+|\d{2,})(?![\d.]*[:!])')
RX_EUR = re.compile(r'(\d{1,3}(?:\.\d{3})+|\d+)(?:,\d+)?\s*(?:€|EUR\b|euros?\b)', re.I)
RX_SECCION = re.compile(r'^(\d{1,2})\. \S')
RX_SUB = re.compile(r'^\d{1,2}\.\d \S')
RX_ACT = re.compile(r'customer|\bcheck\b|per day|operating day|truck|window|street|vending', re.I)

# --- clasificación de las 108 tareas del checklist ES por FILA (criterio F1, ver F1-inventario-delta.md §5) -----
ES_SOLO = [7, 12, 13, 14, 19, 20, 21, 64, 83]      # IAE, RGSEAA, RGPD, LSSI, modelos 390/347/349, Hosply ×2
ADAPTAR = [6, 9, 10, 11, 15, 17, 18, 22, 25, 26, 27, 28, 29, 30, 31, 32, 33, 34, 58, 59, 60, 61, 62, 63, 70, 77, 81,
           82, 95, 110]                             # trámite/proveedor/moneda españoles con equivalente US
MULTICOCINA = list(range(37, 49))                   # 12 equipos por cocina del mundo (sushi, wok, tandoor, trompo…)
# reparto propuesto a las 6 fases del molde EN (fila ES → fase EN); ES_SOLO no se reparte
FASE_EN = {}
for _f in (6, 8, 15, 16, 17, 18, 22, 25, 26, 27, 28, 29, 30):
    FASE_EN[_f] = 1                                 # Business setup (entidad, impuestos, seguros)
for _f in (9, 11, 31, 32, 33, 34, 63):
    FASE_EN[_f] = 2                                 # Kitchen, van & permits
for _f in range(37, 63):
    FASE_EN[_f] = 3                                 # Equipment & suppliers
FASE_EN[10] = 4                                     # Staff (food handler / manager certification)
for _f in list(range(67, 83)) + list(range(86, 102)):
    FASE_EN[_f] = 5                                 # Marketing & sales
for _f in range(104, 124):
    FASE_EN[_f] = 6                                 # First events


def sha(p):
    with open(p, 'rb') as f:
        return hashlib.sha256(f.read()).hexdigest()


def palabras(s):
    return len(re.findall(r"[A-Za-zÀ-ÿ0-9]+(?:[.,'’-][A-Za-zÀ-ÿ0-9]+)*", s))


def es_fx(v):
    return isinstance(v, str) and v.startswith('=')


def libro_xlsx(ruta, detalle=True):
    wb = openpyxl.load_workbook(ruta)
    wv = openpyxl.load_workbook(ruta, data_only=True)
    out = collections.OrderedDict()
    for ws in wb.worksheets:
        fx, txt, w, const, lit = 0, 0, 0, [], []
        rellenos = collections.Counter()
        for row in ws.iter_rows():
            for c in row:
                v = c.value
                if v is None:
                    continue
                if es_fx(v):
                    fx += 1
                    sin_ref = re.sub(r"'[^']+'!|\$?[A-Z]{1,3}\$?\d+", '', v)
                    nums = [n for n in RX_NUM_EN_FX.findall(sin_ref) if n not in ('100',)]
                    if detalle and nums:
                        const.append([c.coordinate, v[:120], wv[ws.title][c.coordinate].value])
                    if '"' in v:
                        lit.append(c.coordinate)
                elif isinstance(v, str):
                    txt += 1
                    w += palabras(v)
                if c.fill is not None and c.fill.fill_type:
                    rellenos[str(c.fill.fgColor.rgb)[-6:]] += 1
        out[ws.title] = collections.OrderedDict([
            ('rango', ws.dimensions), ('formulas', fx), ('textos', txt), ('palabras', w),
            ('dv', len(ws.data_validations.dataValidation)), ('cf', len(ws.conditional_formatting)),
            ('merges', len(ws.merged_cells.ranges)), ('protegida', bool(ws.protection.sheet)),
            ('rellenos', dict(rellenos.most_common(6))), ('fx_con_constantes', const), ('fx_con_texto', lit)])
    return wb, wv, out


def caso_es(wv):
    """Cifras del caso ES leídas de la caché y coherencias internas."""
    r, pr, pe, mix, kpi = (wv['Resumen'], wv['Proyección 3 años'], wv['Break-even'], wv['Mix B2C-B2B'], wv['KPIs'])
    pl = wv['P&L Año 1']
    seg = collections.defaultdict(lambda: [0, 0.0])
    for f in range(5, 15):
        s, t, n = mix[f'A{f}'].value, mix[f'C{f}'].value, mix[f'D{f}'].value
        seg[s][0] += n
        seg[s][1] += t * n
    pond = {s: round(v[1] / v[0], 1) for s, v in seg.items()}
    meses = [pl.cell(7, c).value for c in range(2, 14)]
    return collections.OrderedDict([
        ('inversion_min_rec_prem', [r['B21'].value, r['C21'].value, r['D21'].value]),
        ('inversion_suma_min_tecleada', sum(r[f'B{f}'].value for f in range(6, 21))),
        ('ingresos_a1_a2_a3', [pr['B11'].value, pr['C11'].value, pr['D11'].value]),
        ('neto_a1_a2_a3', [pr['B17'].value, pr['C17'].value, pr['D17'].value]),
        ('margen_neto_a1_a2_a3', [round(pr[f'{c}18'].value, 4) for c in 'BCD']),
        ('eventos_b2c_b2b_a1', [pr['B5'].value, pr['B6'].value]),
        ('ticket_b2c_b2b_proyeccion', [pr['B7'].value, pr['B8'].value]),
        ('eventos_mix', {s: v[0] for s, v in seg.items()}),
        ('ticket_ponderado_mix', pond),
        ('fijos_anuales_PE_B5', pe['B5'].value),
        ('pe_eventos_anio_mes', [round(pe['B9'].value, 2), round(pe['B10'].value, 3)]),
        ('pyg_a1_ingresos_mensuales', meses),
        ('pyg_a1_estacionalidad_pct', [round(m / sum(meses), 4) for m in meses]),
        ('pyg_a1_neto', pl['N19'].value),
        ('kpi_margen_neto_objetivo', kpi['B8'].value),
        ('kpi_eventos_mes', kpi['B5'].value),
    ])


def checklist_es(wb):
    ws = wb.worksheets[0]
    fases, tareas = collections.OrderedDict(), []
    actual = None
    for f in range(5, ws.max_row + 1):
        a, b = ws[f'A{f}'].value, ws[f'B{f}'].value
        if isinstance(a, str) and re.match(r'^\d\. ', a):
            actual = a
            fases[actual] = []
        elif isinstance(b, int):
            fases[actual].append(f)
            tareas.append(f)
    clas = collections.Counter()
    for f in tareas:
        clas['ES_SOLO' if f in ES_SOLO else 'MULTICOCINA' if f in MULTICOCINA else
             'ADAPTAR' if f in ADAPTAR else 'directo'] += 1
    por_fase = collections.OrderedDict()
    for fase, filas in fases.items():
        por_fase[fase] = collections.OrderedDict([
            ('tareas', len(filas)), ('filas', [filas[0], filas[-1]]),
            ('es_solo', [f for f in filas if f in ES_SOLO]),
            ('adaptar', len([f for f in filas if f in ADAPTAR])),
            ('multicocina', len([f for f in filas if f in MULTICOCINA]))])
    reparto = collections.Counter(FASE_EN[f] for f in tareas if f not in ES_SOLO)
    sin_fase = [f for f in tareas if f not in ES_SOLO and f not in FASE_EN]
    cab = [ws.cell(4, c).value for c in range(1, 7)]
    return collections.OrderedDict([
        ('titulo_A2', ws['A2'].value), ('cabecera_fila4', cab), ('total_tareas', len(tareas)),
        ('por_fase', por_fase), ('clasificacion', dict(clas)),
        ('reparto_fases_en', dict(sorted(reparto.items()))), ('sin_fase_en', sin_fase),
        ('palabras_hito_detalle', sum(palabras(str(ws[f'C{f}'].value or '')) + palabras(str(ws[f'D{f}'].value or ''))
                                      for f in tareas))])


def docx_es(ruta):
    d = docx.Document(ruta)
    ps = d.paragraphs
    secciones, actual = collections.OrderedDict(), 'portada'
    secciones[actual] = [0, 0, []]
    for i, p in enumerate(ps):
        t = p.text.strip()
        if RX_SECCION.match(t) and len(t) < 70:
            actual = t
            secciones[actual] = [0, 0, [i]]
            continue
        if t:
            secciones[actual][0] += 1
            secciones[actual][1] += palabras(t)
    tabw = sum(palabras(c.text) for t in d.tables for r in t.rows for c in r.cells)
    eur = [(i, m.group(0)) for i, p in enumerate(ps) for m in RX_EUR.finditer(p.text)]
    return collections.OrderedDict([
        ('parrafos', len(ps)), ('con_texto', sum(1 for p in ps if p.text.strip())),
        ('estilos', dict(collections.Counter(p.style.name for p in ps))),
        ('headings', sum(1 for p in ps if p.style.name.startswith('Heading'))),
        ('secciones', len([s for s in secciones if s != 'portada'])),
        ('subsecciones', sum(1 for p in ps if RX_SUB.match(p.text.strip()))),
        ('palabras_parrafos', sum(palabras(p.text) for p in ps)), ('palabras_tablas', tabw),
        ('por_seccion', {k: {'parrafos': v[0], 'palabras': v[1], 'idx': v[2]} for k, v in secciones.items()}),
        ('tablas', [[len(t.rows), len(t.columns), t.rows[0].cells[0].text[:30]] for t in d.tables]),
        ('importes_eur', eur), ('pagina_emu', [d.sections[0].page_width, d.sections[0].page_height])])


def docx_en(ruta):
    d = docx.Document(ruta)
    ps = d.paragraphs
    return collections.OrderedDict([
        ('parrafos', len(ps)), ('headings_idx', [i for i, p in enumerate(ps) if p.style.name.startswith('Heading')]),
        ('palabras', sum(palabras(p.text) for p in ps)), ('tablas', len(d.tables))])


def libro_en(repo, corto):
    carpeta, pref = EN[corto]
    base = os.path.join(repo, DL, carpeta)
    plan = os.path.join(base, f'{pref}-financial-projections.xlsx')
    wb, wv, est = libro_xlsx(plan, detalle=False)
    def hoja(n):                                    # REST numera las pestañas («1. Startup Costs»)
        return next(wb[t] for t in wb.sheetnames if n in (t, re.sub(r'^[1-9]\. ', '', t)))
    s, inv, stf, pl = hoja('0. Assumptions'), hoja('Startup Costs'), hoja('Staffing'), hoja('3-Year P&L')
    tecleadas = [f for f in range(1, inv.max_row + 1)
                 if isinstance(inv[f'B{f}'].value, (int, float)) and not es_fx(inv[f'B{f}'].value)]
    puestos = [f for f in range(5, 20) if isinstance(stf[f'D{f}'].value, (int, float))]
    tot_staff = next((f for f in range(5, 20) if str(stf[f'A{f}'].value or '').startswith('TOTAL')), None)
    a_fijos = next(f for f in range(1, 40) if str(pl[f'A{f}'].value or '').startswith('FIXED COSTS'))
    a_tot = next(f for f in range(1, 60) if str(pl[f'A{f}'].value or '').startswith('TOTAL FIXED'))
    fijos = [str(pl[f'A{f}'].value) for f in range(a_fijos + 1, a_tot)]
    cf = hoja('12-Month Cash Flow')
    ck_ruta = [os.path.join(base, n) for n in os.listdir(base) if n.endswith('checklist.xlsx')][0]
    ck = openpyxl.load_workbook(ck_ruta)
    tareas = []
    for ws in ck.worksheets:
        if ws.title.startswith('Phase'):
            tareas.append(sum(1 for f in range(5, ws.max_row + 1)
                              if ws[f'A{f}'].value in ('☐', '✓', 'N/A') and ws[f'B{f}'].value))
    dx = [n for n in os.listdir(base) if n.endswith('.docx')][0]
    return collections.OrderedDict([
        ('formulas', sum(v['formulas'] for v in est.values())),
        ('hojas', {k: [v['rango'], v['formulas']] for k, v in est.items()}),
        ('B4_B5_B6', [s['A4'].value, s['A5'].value, s['A6'].value]),
        ('B24', [s['A24'].value, s['B24'].value]),
        ('activacion_B4_B5_B6', [s['B4'].value, s['B5'].value, s['B6'].value]),
        ('startup_tecleadas', len(tecleadas)), ('startup_filas', [tecleadas[0], tecleadas[-1]] if tecleadas else []),
        ('startup_rotulos', [str(inv[f'A{f}'].value) for f in tecleadas]),
        ('staffing_puestos', len(puestos)), ('staffing_total_fila', tot_staff),
        ('staffing_fte_total', stf[f'C{tot_staff}'].value if tot_staff else None),
        ('pyg_fijos', fijos), ('estacionalidad_A5', cf['A5'].value),
        ('estacionalidad', [cf.cell(5, c).value for c in range(2, 14)]),
        ('checklist', os.path.basename(ck_ruta)), ('checklist_tareas_por_fase', tareas),
        ('docx', dict(docx_en(os.path.join(base, dx)), fichero=dx))])


def drivers_ft(repo):
    """Fórmulas del FT EN que leen cada driver de actividad (directo en «0. Assumptions» B4/B5/B6 o vía las filas
    6/7/8 del P&L que los reexponen); celdas de texto que nombran la actividad del food truck."""
    ruta = os.path.join(repo, DL, 'food-truck-business-plan', 'food-truck-financial-projections.xlsx')
    wb = openpyxl.load_workbook(ruta)
    rx_sup = re.compile(r"'0\. Assumptions'!\$?B\$?([456])(?!\d)")
    rx_pl_fuera = re.compile(r"'3-Year P&L'!\$?[BCD]\$?([678])(?!\d)")
    rx_pl_dentro = re.compile(r"(?<![A-Z'!$])\$?[BCD]\$?([678])(?!\d)")
    nombre = {'4': 'B4_actividad', '5': 'B5_ticket', '6': 'B6_dias', '6p': 'PL6_actividad', '7p': 'PL7_ticket',
              '8p': 'PL8_dias'}
    uso = collections.defaultdict(set)
    textos_act = []
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                v = c.value
                if es_fx(v):
                    for m in rx_sup.findall(v):
                        uso[nombre[m]].add(f'{ws.title}!{c.coordinate}')
                    rx = rx_pl_dentro if ws.title == '3-Year P&L' else rx_pl_fuera
                    for m in rx.findall(re.sub(r"'[^']+'!\$?[A-Z]+\$?\d+", '', v) if ws.title == '3-Year P&L'
                                        else v):
                        uso[nombre[m + 'p']].add(f'{ws.title}!{c.coordinate}')
                elif isinstance(v, str) and RX_ACT.search(v):
                    textos_act.append(f'{ws.title}!{c.coordinate}')
    return {'uso': {k: sorted(v) for k, v in sorted(uso.items())},
            'n_uso': {k: len(v) for k, v in sorted(uso.items())},
            'n_textos_actividad': len(textos_act),
            'textos_actividad_por_hoja': dict(collections.Counter(t.split('!')[0] for t in textos_act))}


def ficha(repo):
    t = open(os.path.join(repo, FICHA_ES), encoding='utf-8').read()
    claves = [r'5\.500', r'11\.500', r'35\.000', r'14 eventos/mes', r'60\+ pp', r'110 hitos', r'10 ejemplos validados',
              r'96 proveedores', r'priceOld', r'73 % de descuento']
    return {k.replace('\\', ''): len(re.findall(k, t)) for k in claves}


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--repo', default=os.path.abspath(os.path.join(AQUI, '..', '..', '..')))
    ap.add_argument('--salida', default=os.path.join(AQUI, 'inventario_cat.json'))
    a = ap.parse_args()
    base = os.path.join(a.repo, DL, ES_DIR)
    res = collections.OrderedDict()
    res['sha256'] = {n: sha(os.path.join(base, n)) for n in sorted(os.listdir(base))}
    res['en_alcance'] = list(ES_ALCANCE.values())
    _, wv, res['CATP'] = libro_xlsx(os.path.join(base, ES_ALCANCE['CATP']))
    res['caso_es'] = caso_es(wv)
    _, _, res['CATPR'] = libro_xlsx(os.path.join(base, ES_ALCANCE['CATPR']))
    wbc, _, res['CATC_libro'] = libro_xlsx(os.path.join(base, ES_ALCANCE['CATC']), detalle=False)
    res['CATC'] = checklist_es(wbc)
    res['CATD'] = docx_es(os.path.join(base, ES_ALCANCE['CATD']))
    res['docx_ft_es'] = docx_en(os.path.join(a.repo, DL, *DOCX_FT_ES))
    res['en'] = collections.OrderedDict((k, libro_en(a.repo, k)) for k in EN)
    res['drivers_ft'] = drivers_ft(a.repo)
    res['ficha_es'] = ficha(a.repo)
    with open(a.salida, 'w', encoding='utf-8') as f:
        json.dump(res, f, ensure_ascii=False, indent=1, default=str)
    print('OK', a.salida, hashlib.sha256(open(a.salida, 'rb').read()).hexdigest()[:12])


if __name__ == '__main__':
    main()
