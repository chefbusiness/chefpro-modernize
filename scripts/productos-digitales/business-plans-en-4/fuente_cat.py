#!/usr/bin/env python3
"""
fuente_cat.py — material de CONTENIDO para los redactores: el checklist y el plan de negocio del catering ES (línea B,
`plan-catering-tematico-eventos`) en `fuente_es_cat.json`. El circuito de texto corre sobre el FOOD TRUCK ES (las
ranuras); el contenido de catering sale de aquí y de research (D71, D77, D78). Corre en el VPS (openpyxl +
python-docx). Sesión Claude Code, 11-oct-2026 (F2, paso 1).

    python3 bp4.py fuente
"""
import datetime
import hashlib
import json
import os
from collections import OrderedDict

import docx
import openpyxl

AQUI = os.path.dirname(os.path.abspath(__file__))
REPO = os.path.abspath(os.path.join(AQUI, '..', '..', '..'))
DIR = os.path.join(REPO, 'astro-site', 'public', 'dl', 'plan-catering-tematico-eventos')
CHECK = 'checklist-apertura-catering-tematico-eventos.xlsx'
DOCX = 'plan-de-negocio-catering-tematico-eventos.docx'
# F1-inventario-delta §5: filas sin sentido en EE. UU. y filas del equipo multicocina (fuera del EN)
SOLO_ES = {7, 12, 13, 14, 19, 20, 21, 64, 83}
MULTICOCINA = set(range(37, 49))
FASE_EN = {'1': 'P1 Business Setup (+ P2 / P4)', '2': 'P1 / P2 Kitchen, Van & Permits', '3': 'P3 Equipment (+ P2 van)',
           '4': 'P5 Marketing', '5': 'P5 Marketing', '6': 'P6 First 90 Days'}


def sha(p):
    return hashlib.sha256(open(p, 'rb').read()).hexdigest()


def checklist():
    p = os.path.join(DIR, CHECK)
    ws = openpyxl.load_workbook(p, data_only=True).worksheets[0]
    fases, fase = OrderedDict(), None
    for r in range(5, ws.max_row + 1):
        a, b, c, d, e = (ws.cell(r, k).value for k in range(1, 6))
        if isinstance(a, str) and a.strip() and not c:
            fase = a.strip()
            fases[fase] = []
            continue
        if fase is None or not (isinstance(c, str) and c.strip()):
            continue
        marca = 'SOLO_ES (fuera)' if r in SOLO_ES else ('MULTICOCINA (fuera)' if r in MULTICOCINA else '')
        x = OrderedDict([('fila', r), ('hito', c.strip()), ('detalle', (d or '').strip() if isinstance(d, str) else d),
                         ('plazo', e)])
        if marca:
            x['marca'] = marca
        fases[fase].append(x)
    return p, fases


def plan_docx():
    p = os.path.join(DIR, DOCX)
    d = docx.Document(p)
    pars = [OrderedDict([('i', i), ('estilo', x.style.name), ('texto', x.text)]) for i, x in enumerate(d.paragraphs)
            if x.text.strip()]
    tablas = [[[c.text for c in row.cells] for row in t.rows] for t in d.tables]
    return p, pars, tablas


def main():
    pc, fases = checklist()
    pd, pars, tablas = plan_docx()
    for k in list(fases):
        n = k.split('.')[0].strip()
        fases[k] = OrderedDict([('fase_en_sugerida', FASE_EN.get(n, '')), ('tareas', fases[k])])
    out = OrderedDict([
        ('meta', OrderedDict([
            ('generado', datetime.datetime.now(datetime.timezone.utc).isoformat(timespec='seconds')),
            ('script', 'scripts/productos-digitales/business-plans-en-4/fuente_cat.py'),
            ('uso', 'CONTENIDO de catering para los redactores (D71, D77, D78). NO se traduce tal cual: se condensa y '
                    'se adapta a EE. UU. con research (F1-research-us-delta.md). Las cifras del ES NO se usan (D70): '
                    'las del caso salen del libro EN (tokens). Fuera: tareas marcadas SOLO_ES y MULTICOCINA, «España», '
                    '«ChefBusiness», «96 proveedores», «12 cocinas del mundo», importes en €'),
            ('fuentes', OrderedDict([(os.path.relpath(pc, REPO), sha(pc)), (os.path.relpath(pd, REPO), sha(pd))])),
        ])),
        ('checklist_es', fases),
        ('plan_docx_es', OrderedDict([('parrafos', pars), ('tablas', tablas)])),
    ])
    dest = os.path.join(AQUI, 'fuente_es_cat.json')
    with open(dest, 'w', encoding='utf-8') as fh:
        json.dump(out, fh, ensure_ascii=False, indent=1)
        fh.write('\n')
    n = sum(len(f['tareas']) for f in fases.values())
    fuera = sum(1 for f in fases.values() for t in f['tareas'] if 'marca' in t)
    print('fuente_es_cat.json: %d fases, %d tareas (%d fuera), %d párrafos, %d tablas' % (
        len(fases), n, fuera, len(pars), len(tablas)))
    return 0


if __name__ == '__main__':
    raise SystemExit(main())
