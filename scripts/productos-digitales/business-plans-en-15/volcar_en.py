#!/usr/bin/env python3
"""
volcar_en.py — vuelca a TEXTO los 3 entregables EN del Burger Restaurant Business Plan Kit
(`astro-site/public/dl/burger-restaurant-business-plan/`) para la revisión adversarial final, que se hace en el Mac sin abrir
xlsx/docx (como los `dump-en/` de la pizza, que no se versionaron). Corre en el VPS (openpyxl + python-docx) y escribe
en `business-plans-en-15/revision/` (o --salida):

  burger-restaurant-financial-projections.txt   por hoja: «CELDA ⇥ valor» (fórmulas: «= valor en caché ‖ fórmula»), DV,
                                              docProps y la configuración de página
  burger-restaurant-startup-checklist.txt       ídem
  burger-restaurant-business-plan.txt           «[i] (estilo) texto» de cada párrafo con texto, docProps y página

Sesión Claude Code, 29-oct-2026 (F2-EN, paso 4).

    /root/venv-guias/bin/python volcar_en.py [--dir DL] [--salida DIR]
"""
import argparse
import os

import docx
import openpyxl

AQUI = os.path.dirname(os.path.abspath(__file__))
REPO = os.path.abspath(os.path.join(AQUI, '..', '..', '..'))
DL = os.path.join(REPO, 'astro-site', 'public', 'dl', 'burger-restaurant-business-plan')


def _v(x):
    if isinstance(x, float):
        return repr(round(x, 6))
    return str(x).replace('\n', ' ⏎ ')


def volcar_xlsx(p):
    wf = openpyxl.load_workbook(p)
    wv = openpyxl.load_workbook(p, data_only=True)
    out = ['# %s' % os.path.basename(p)]
    pr = wf.properties
    out.append('docProps · title: %s · subject: %s · keywords: %s · description: %s' % (
        pr.title, pr.subject, pr.keywords, pr.description))
    for ws in wf.worksheets:
        wsv = wv[ws.title]
        ps = ws.page_setup
        out += ['', '## %s  (papel %s, orientación %s, protegida %s)' % (ws.title, ps.paperSize, ps.orientation,
                                                                       ws.protection.sheet)]
        for row in ws.iter_rows():
            for c in row:
                if c.value is None:
                    continue
                if isinstance(c.value, str) and c.value.startswith('='):
                    out.append('%s\t= %s ‖ %s' % (c.coordinate, _v(wsv[c.coordinate].value), c.value))
                else:
                    out.append('%s\t%s' % (c.coordinate, _v(c.value)))
        for dv in ws.data_validations.dataValidation:
            out.append('DV %s\t%s · %s · error: %s · ayuda: %s' % (dv.sqref, dv.type, dv.formula1, dv.error, dv.prompt))
    return '\n'.join(out) + '\n'


def volcar_docx(p):
    d = docx.Document(p)
    s = d.sections[0]
    cp = d.core_properties
    out = ['# %s' % os.path.basename(p),
           'core · title: %s · subject: %s · keywords: %s · author: %s' % (cp.title, cp.subject, cp.keywords, cp.author),
           'página %.2f × %.2f in · %d párrafos · %d palabras' % (
               s.page_width.inches, s.page_height.inches, len(d.paragraphs),
               sum(len(x.text.split()) for x in d.paragraphs)), '']
    for i, x in enumerate(d.paragraphs):
        if x.text.strip():
            out.append('[%d] (%s) %s' % (i, x.style.name, x.text.replace('\n', ' ⏎ ')))
    return '\n'.join(out) + '\n'


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--dir', default=DL)
    ap.add_argument('--salida', default=os.path.join(AQUI, 'revision'))
    a = ap.parse_args()
    os.makedirs(a.salida, exist_ok=True)
    for f in sorted(os.listdir(a.dir)):
        p = os.path.join(a.dir, f)
        if f.endswith('.xlsx'):
            t = volcar_xlsx(p)
        elif f.endswith('.docx'):
            t = volcar_docx(p)
        else:
            continue
        q = os.path.join(a.salida, os.path.splitext(f)[0] + '.txt')
        with open(q, 'w', encoding='utf-8') as fh:
            fh.write(t)
        print('%s → %s (%d líneas)' % (f, os.path.relpath(q, REPO), t.count('\n')))


if __name__ == '__main__':
    main()
