#!/usr/bin/env python3
"""
calculadora_bur.py — Smash Burger Patty, Fries & Costing Calculator, calculadora del Burger Restaurant Business Plan Kit (EN, US).
Es la edición EE. UU. de `../plan-hamburgueseria/calculadora_ham_es.py`: MISMAS hojas (`Beef & Patty` · `Fries & Oil` · `Costing` ·
`Instructions`), filas, fórmulas, helpers, `inject_cache` y autotest (pycel + caché = `calcular()`, 3 juegos), con el patrón de
`../business-plans-en-14/calculadora_sus.py`. Cambios EN:

  (1) carne, patatas y aceite en libras ($/lb) y porciones en ONZAS, con la constante gris «16 oz per lb» (el gate prohíbe
      números tecleados en fórmulas salvo 0/1); aceite por PESO (lb: las freidoras de EE. UU. se miden en lb de aceite);
      «60 minutes per hour»; °F solo en el texto de las instrucciones;
  (2) coste por lb útil = precio ÷ útil ÷ (1 − merma de cierre); mano de obra del cocinero en $/h;
  (3) cocción 155 °F (68 °C) 17 s (FDA Food Code 2022, 3-401.11(A)(2)); consumer advisory 3-603.11 solo si se sirve poco hecha a
      petición; registro de picado como buena práctica (nunca se afirma que sea obligatorio); TPM sin límite federal;
  (4) hoja `Costing`: sales tax en carta = 0 %, comisión del marketplace 25 %; queso en $ por loncha (un solo dato en lugar del
      par €/kg + g por loncha); margen dine-in · take-out/pickup · marketplace;
  (5) US Letter, formatos sin símbolo de moneda (los rótulos llevan «($)»), docProps en inglés, versión del libro EN.

Valores que otros scripts leen de la CACHÉ del xlsx (dict `EXPUESTOS`, función `valores_expuestos(ruta)`).
Corre en el VPS (openpyxl + pycel):

    /root/venv-guias/bin/python calculadora_bur.py --salida <dir>             # genera + inject_cache + autotest
    /root/venv-guias/bin/python calculadora_bur.py --salida <dir> --autotest

Sesión Claude Code, 31-oct-2026.
"""
import argparse
import math
import os
import re
import sys
from collections import OrderedDict

import openpyxl
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill, Protection
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

AQUI = os.path.dirname(os.path.abspath(__file__))
PD = os.path.dirname(AQUI)
os.environ.setdefault('BP_CONJUNTO', 'bur')
sys.path.insert(0, os.path.join(PD, 'planes-v2_0'))
sys.path.insert(0, os.path.join(PD, 'business-plans-en'))
sys.path.insert(0, PD)
sys.path.insert(0, AQUI)
import motor                                                                       # noqa: E402  (solo gate_formatos)
import mapas as M                                                                  # noqa: E402  (detectores y textos fijos)

_PL = M.PLANES['bur']
SLUG = _PL['slug']
PRODUCTO = _PL['producto']
URL = _PL['url']
FICHERO = 'smash-burger-patty-costing-calculator.xlsx'
TITULO = 'Smash Burger Patty, Fries & Costing Calculator (Beef & Patty, Fries & Oil and Cost per Burger)'
KEYWORDS = _PL['keywords']
VERSION, MES = M.VERSION, M.MES_VERSION
VERSION_LINEA = 'Version %s · %s · %s · info@aichef.pro' % (VERSION, MES, URL)
FOOTER = M.FOOTER_EN
HC, HP, HE, HI = 'Beef & Patty', 'Fries & Oil', 'Costing', 'Instructions'
VERDE = PatternFill('solid', fgColor='E8F5E9')
GRIS = PatternFill('solid', fgColor='F2F2F2')
ROJO = PatternFill('solid', bgColor='F8CBAD', fgColor='F8CBAD')
AMBAR = PatternFill('solid', bgColor='FFE699', fgColor='FFE699')
OK = PatternFill('solid', bgColor='C6E0B4', fgColor='C6E0B4')
FUENTE = 'Calibri'
USD, ENT, PCT, DEC, OZF, USD3 = '#,##0.00', '#,##0', '0.0%', '#,##0.0', '#,##0.00', '#,##0.000'
FC_OK, FC_VIG, FC_ALTO = 'On target', 'Watch the cost', 'Above target'
SEM_PICAR, SEM_BLEND = 'Grinding in-house pays off', 'Buy the blend ground'
SEM_FRE, SEM_CONG = 'Fresh-cut pays off', 'Frozen pays off'
EST = ' Kit estimate, editable.'
TAX_ROT = 'Sales tax included in your menu prices (%)'

# ---------------------------------------------------------------------------------------------------------------------
# Datos (F1-research-us-delta §4-§8 y SPEC-calculadora). Importes antes de sales tax. Todo «Kit estimate (Oct 2026)», editable.
# ---------------------------------------------------------------------------------------------------------------------
NPE = 4
PC = [get_column_letter(2 + i) for i in range(NPE)]                                # B..E
#: hoja Beef & Patty, bloque A: (nombre, precio $/lb, útil, minutos de recortar y picar por lb, % de grasa)
CARNES = [
    ('80/20 ground chuck blend (bought ground)', 5.25, 1.00, 0.0, 0.20),
    ('Whole chuck for house grinding', 4.50, 0.90, 3.0, 0.20),
    ('Pre-portioned balls (bought, 3 oz)', 6.25, 1.00, 0.0, 0.20),
    ('Other beef (type yours)', 5.75, 1.00, 0.0, 0.20),
]
P_ROW = OrderedDict([('pr', 6), ('ut', 7), ('tl', 8), ('gr', 9)])
P_NOTAS = OrderedDict([
    ('pr', 'No wholesale price for 80/20 was read in Oct 2026: BLS/FRED ground beef (100% beef, retail) came in at $6.79/lb in Sep 2026 '
           'and retail 80/20 runs about $5.69-6.93/lb. The kit puts 5.25 $/lb for 80/20 chuck from a distributor, 4.50 for whole chuck '
           'and 6.25 for pre-portioned balls (no source: ask your supplier). Kit estimate (Oct 2026): replace with your '
           "distributor's price."),
    ('ut', 'Share of the beef you buy that is left after trimming and grinding. 100% if you buy it ground; 90% for whole chuck (no source: '
           'weigh before and after).' + EST),
    ('tl', 'Minutes to trim and grind 1 lb you buy. No source: time 3 lb and use the average. Only whole chuck carries labor.' + EST),
    ('gr', 'Informational: it changes no figure. The blend of the plan runs about 20% fat.' + EST),
])
OPC = [  # desplegable informativo: (opción, recordatorio)
    ('I buy ground beef',
     "Keep raw ground beef at 41 °F (5 °C) or below and store it below ready-to-eat food (FDA Food Code 2022, 3-501.16 and 3-302.11). "
     "File the supplier's establishment number, lot and pack date of every delivery."),
    ('I grind in-house',
     "Keep a house-grind log as good practice and for recalls: supplier establishment numbers and lots, the date and time of every grind "
     'and of every cleaning and sanitizing of the grinder. The FSIS rule (9 CFR 320.1(b)(4)) covers retail stores and official '
     'establishments that grind raw beef: ask your health department whether it applies to you.'),
    ('I buy pre-portioned balls',
     'Keep the spec sheet and the label with the lot and date of every case, and hold the balls at 41 °F (5 °C) or below until they go '
     'on the flat-top.'),
]
#: bloque B (entradas comunes de la hoja Beef & Patty): (clave, rótulo, formato, valor, constante, nota)
ENT_C = [
    ('m', 'End-of-day shrink (% of the balls, buns and sauces thrown out)', PCT, 0.03, False,
     'Balls left at close that cannot be kept, plus buns and sauces thrown out.' + EST),
    ('gb', 'Ounces per ball (oz)', OZF, 3.0, False, '2-4 oz; 3 oz is the case of the plan: weigh 10 balls and use the average.' + EST),
    ('mc', 'Cook loss on the flat-top (% of the ball weight)', PCT, 0.25, False,
     'No source: weigh 10 balls before and after the flat-top and use the average. It only changes the ounces served.' + EST),
    ('hh', 'Cook cost per hour ($/h)', USD, 18.70, False,
     'A grill cook at about $17/h + 10% employer payroll taxes ≈ $18.70/h.' + EST),
    ('bc', 'Balls per cover', '#,##0.00', 1.6, False, 'Plan criterion: measure it.' + EST),
    ('cd', 'Covers per day', ENT, 195, False, 'The case of the business plan. Replace it with yours.' + EST),
    ('cong', 'Do you buy ground beef or grind in-house? (dropdown)', 'General', OPC[0][0], False,
     'Informational: pick your case and read the reminder below. It changes no figure.'),
    ('OZ', 'Ounces per lb (constant: 16 oz per lb)', ENT, 16, True, 'Constant: 1 lb = 16 oz. Not editable.'),
    ('MH', 'Minutes per hour (constant: 60 minutes per hour)', ENT, 60, True, 'Constant: 1 hour = 60 minutes. Not editable.'),
]
PE_M0 = 12
M_C = OrderedDict((k[0], PE_M0 + i) for i, k in enumerate(ENT_C))                  # 12..20
PE_RC = PE_M0 + len(ENT_C) + 1                                                     # 22: cabecera del bloque C
SAL_C = [
    ('cku', 'COST PER USABLE LB ($)', USD, '{pr}/{ut}/(1-{m})', 'Price ÷ usable ÷ (1 − end-of-day shrink).'),
    ('cku_mo', 'Cost per usable lb with the labor of trimming and grinding ($)', USD, '({pr}+{tl}/{MH}*{hh})/{ut}/(1-{m})',
     '(Price + minutes ÷ 60 × cost per hour) ÷ usable ÷ (1 − end-of-day shrink). It feeds the costing.'),
    ('c_bola', 'Cost of the beef of one ball ($)', USD3, '{cku_mo}*{gb}/{OZ}', 'Cost per usable lb with labor × ounces ÷ 16.'),
    ('c_dos', 'Cost of the beef of a double (2 balls, $)', USD3, '{c_bola}+{c_bola}', 'Two balls.'),
    ('c_tres', 'Cost of the beef of a triple (3 balls, $)', USD3, '{c_bola}+{c_bola}+{c_bola}', 'Three balls.'),
]
S_C = OrderedDict((k[0], PE_RC + 1 + i) for i, k in enumerate(SAL_C))             # 23..27
PE_RD = PE_RC + len(SAL_C) + 2                                                     # 29: cabecera del bloque D
SAL_D = [
    ('bkg', 'BALLS PER LB, RAW (balls)', DEC, '{OZ}/{gb}', '16 ÷ ounces per ball. It is the figure on the supplier spec sheet.'),
    ('bkn', 'Balls per lb counting the end-of-day shrink (balls)', DEC, '{OZ}*(1-{m})/{gb}', '16 × (1 − end-of-day shrink) ÷ ounces per ball.'),
    ('gser', 'Ounces of beef you serve after the flat-top (oz)', OZF, '{gb}*(1-{mc})', 'Ounces per ball × (1 − cook loss).'),
    ('bdia', 'Balls you sell per day (balls)', ENT, '{cd}*{bc}', 'Covers per day × balls per cover.'),
    ('carne_dia', 'LB OF BEEF PER DAY', '#,##0.0', '{bdia}*{gb}/{OZ}',
     'Balls per day × ounces per ball ÷ 16. It sizes the daily order and the day\'s grinding if you grind in-house.'),
    ('lab_kg', 'Labor of trimming and grinding per usable lb of whole chuck ($)', USD, '{cku_mo1}-{cku1}',
     'Cost per usable lb of whole chuck with labor − without it.'),
    ('X', 'GRINDING IN-HOUSE PAYS OFF AS LONG AS TRIMMING AND GRINDING COST LESS THAN ($ per usable lb)', USD, '{cku_mo0}-{cku1}',
     'Cost per usable lb of the blend bought ground (with labor) − cost per usable lb of whole chuck (without labor).'),
    ('sem', 'Verdict with your labor', 'General', 'IF({lab_kg}<{X},"%s","%s")' % (SEM_PICAR, SEM_BLEND),
     'Grinding pays off if trimming and grinding cost less than that difference.'),
    ('rem', 'Reminder for your case', 'General', 'INDEX({rem_r},MATCH({cong},{opc_r},0))', 'Text from the list below.'),
]
S_D = OrderedDict((k[0], PE_RD + 1 + i) for i, k in enumerate(SAL_D))             # 30..38
PE_NOTA = PE_RD + len(SAL_D) + 1                                                   # 39: nota fija bajo los resultados
PE_RO = PE_RD + len(SAL_D) + 2                                                     # 40: cabecera de la lista
PE_O0 = PE_RO + 1                                                                  # 41..43
PE_O1 = PE_O0 + len(OPC) - 1
NOTA_COCCION = ('FIXED NOTE. Cook every smash burger to 155 °F (68 °C) for 17 seconds (FDA Food Code 2022, 3-401.11(A)(2)) and probe one '
                'burger per batch: the crust is not a thermometer.')

#: hoja Fries & Oil: (clave, rótulo, formato, valor, constante, nota)
ENT_A = [
    ('pc', 'Frozen fries ($/lb, before tax)', USD, 1.16, False,
     'Saint-Arneault 1/2" crinkle, 6 × 5 lb $34.68; 3/8" crinkle $1.11/lb; webstaurantstore.com, Oct 2026.' + EST),
    ('gc', 'Ounces of frozen fries per portion (oz)', OZF, 5.0, False, 'Weigh 10 portions and use the average.' + EST),
    ('ac', 'Oil that stays in a frozen portion (% of its weight)', PCT, 0.05, False, 'What the fries absorb while they fry.' + EST),
    ('pf', 'Fresh potatoes, russet ($/lb, before tax)', USD, 0.55, False, 'No price read: ask your produce supplier.' + EST),
    ('gf', 'Ounces of fresh-cut fries per portion (oz)', OZF, 5.0, False, 'Weigh 10 portions and use the average.' + EST),
    ('uf', 'Usable share of fresh potatoes after peeling and cutting (%)', PCT, 0.75, False,
     'No source: weigh 10 lb before and after.' + EST),
    ('af', 'Oil that stays in a fresh portion (% of its weight)', PCT, 0.10, False, 'Fresh-cut fries absorb more oil than frozen.' + EST),
    ('tf', 'Minutes to peel, cut and soak 1 lb of fresh potatoes (min)', DEC, 3.0, False,
     'No source: time 3 lb and use the average.' + EST),
    ('hp', 'Cook cost per hour ($/h)', USD, 18.70, False, 'Same as the Beef & Patty sheet.' + EST),
    ('lc', 'Lb of oil per fryer (lb)', DEC, 40.0, False, 'A 40-50 lb fryer.' + EST),
    ('nc', 'Fryers for fries', ENT, 2, False, 'The third fryer is for the breaded chicken: cross-contact.' + EST),
    ('pa', 'Frying oil ($/lb, before tax)', USD, 1.40, False,
     'A 35 lb jug or JIB runs about $1.20-1.60/lb; no price read.' + EST),
    ('rcm', 'Portions between oil changes', ENT, 1500, False,
     'Set it by your TPM readings: there is no federal limit; filter nightly.' + EST),
    ('rpc', 'Portions of fries per cover', '#,##0.00', 0.8, False, 'Plan criterion: measure it.' + EST),
    ('dm', 'Days open per month', ENT, 30, False, 'The plan opens 360 days a year.' + EST),
    ('OZ', 'Ounces per lb (constant: 16 oz per lb)', ENT, 16, True, 'Constant: 1 lb = 16 oz. Not editable.'),
    ('MH', 'Minutes per hour (constant: 60 minutes per hour)', ENT, 60, True, 'Constant: 1 hour = 60 minutes. Not editable.'),
]
SAL_A = [
    ('cd_c', 'Covers per day (from the Beef & Patty sheet)', ENT, '{cd_carne}', 'Type it on the Beef & Patty sheet.'),
    ('aceite_cambio', 'Oil cost of one full change ($)', USD, '{lc}*{nc}*{pa}', 'Lb of oil per fryer × fryers × $/lb.'),
    ('ace_cong', 'OIL PER FROZEN PORTION ($)', USD3, '{aceite_cambio}/{rcm}+{gc}*{ac}/{OZ}*{pa}',
     'Oil change ÷ portions between changes + oil that stays in the portion.'),
    ('pat_cong', 'Fries of a frozen portion ($)', USD3, '{gc}*{pc}/{OZ}', 'Ounces × $/lb ÷ 16.'),
    ('c_cong', 'COST PER FROZEN PORTION ($)', USD, '{pat_cong}+{ace_cong}', 'Fries + oil.'),
    ('kg_bruto', 'Lb of unpeeled potatoes per fresh portion (lb)', USD3, '{gf}/{uf}/{OZ}', 'Ounces cut ÷ usable share ÷ 16.'),
    ('pat_fre', 'Fresh potatoes of a portion ($)', USD3, '{kg_bruto}*{pf}', 'Lb unpeeled × $/lb.'),
    ('ace_fre', 'Oil per fresh portion ($)', USD3, '{aceite_cambio}/{rcm}+{gf}*{af}/{OZ}*{pa}',
     'Oil change ÷ portions between changes + oil that stays in the portion.'),
    ('c_fre', 'COST PER FRESH PORTION, WITHOUT LABOR ($)', USD, '{pat_fre}+{ace_fre}', 'Potatoes + oil.'),
    ('mo_fre', 'Labor of peeling, cutting and soaking a fresh portion ($)', USD, '{tf}/{MH}*{hp}*{kg_bruto}',
     'Minutes per lb ÷ 60 × cost per hour × lb unpeeled.'),
    ('c_fre_mo', 'Cost per fresh portion with labor ($)', USD, '{c_fre}+{mo_fre}', 'It feeds the costing if you pick fresh-cut.'),
    ('mo_kg', 'Labor per lb of unpeeled fresh potatoes ($)', USD, '{tf}/{MH}*{hp}', 'Minutes per lb ÷ 60 × cost per hour.'),
    ('X', 'FRESH-CUT PAYS OFF AS LONG AS PEELING AND CUTTING COST LESS THAN ($ per lb unpeeled)', USD, '({c_cong}-{c_fre})/{kg_bruto}',
     '(Cost of the frozen portion − cost of the fresh portion without labor) ÷ lb unpeeled per portion.'),
    ('sem', 'Verdict with your labor', 'General', 'IF({mo_kg}<{X},"%s","%s")' % (SEM_FRE, SEM_CONG),
     'Fresh-cut pays off if peeling and cutting cost less than that difference.'),
    ('rac_dia', 'PORTIONS OF FRIES PER DAY', ENT, '{cd_c}*{rpc}', 'Covers per day × portions per cover.'),
    ('kg_cong_dia', 'Lb of frozen fries per day (lb)', '#,##0.0', '{rac_dia}*{gc}/{OZ}', 'Portions per day × ounces ÷ 16.'),
    ('kg_fre_dia', 'Lb of fresh potatoes, unpeeled, per day (lb)', '#,##0.0', '{rac_dia}*{kg_bruto}',
     'Portions per day × lb unpeeled per portion.'),
    ('l_mes', 'LB OF OIL PER MONTH (frozen fries)', ENT, '{rac_dia}*{dm}*({lc}*{nc}/{rcm}+{gc}*{ac}/{OZ})',
     'Portions per month × (lb of one change ÷ portions between changes + oil that stays in the portion).'),
]
NOTA_FRITURA = ("FIXED NOTE. Fry at 350-355 °F; filter the oil every night and change it by your TPM reading or the manufacturer's guidance; "
                'used oil goes to a licensed renderer. Acrylamide has no federal limit for restaurants; in California, ask your attorney '
                'about the Prop 65 warning for fried potatoes.')
M_A, M_RA, S_A = None, None, None


def _filas(ent, sal, h0=5):
    m = OrderedDict((k[0], h0 + 1 + i) for i, k in enumerate(ent))
    ra = h0 + 1 + len(ent) + 1
    s = OrderedDict((k[0], ra + 1 + i) for i, k in enumerate(sal))
    return m, ra, s


M_A, M_RA, S_A = _filas(ENT_A, SAL_A)
PA_NOTA = max(S_A.values()) + 1                                                    # nota fija bajo los resultados

#: hoja Costing: datos generales
ESC_IN = [
    ('tax', TAX_ROT, PCT, 0.0,
     'US menus usually exclude sales tax: leave 0%. Type your rate only if your prices include it. UK: VAT is 20% dine-in and on hot '
     'take-out food.'),
    ('obj', 'Target food cost (% of the price before tax)', PCT, 0.30,
     "A common target for a burger menu; the plan's 34% is the real usage of the whole food menu." + EST),
    ('tol', 'Watch band above the target (points)', PCT, 0.03, 'Up to this band above the target, the traffic light says "Watch the cost".'),
    ('sel', 'Beef of the costing (1 to 4: the column of the Beef & Patty sheet)', '0', 1,
     '1 = 80/20 blend bought ground, 2 = whole chuck for house grinding, 3 = pre-portioned balls, 4 = free.'),
    ('tp', 'Fries of the costing (1 frozen, 2 fresh-cut with labor)', '0', 1, '1 = frozen portion, 2 = fresh-cut portion with the labor of peeling and cutting.'),
    ('pan', 'Potato bun ($ per bun, before tax)', USD, 0.55,
     "Martin's Sandwich Potato Rolls 2.1 oz, 60 for $33.02 = $0.55, webstaurantstore.com, Oct 2026." + EST),
    ('qs', 'American cheese ($ per slice, before tax)', USD, 0.14,
     'Kraft 640 ct $87.22 ≈ $0.14; Schreiber 480 ct $58.90 ≈ $0.12; webstaurantstore.com. The cost of the cheese is slices × $ per slice.' + EST),
    ('cl', 'To-go packaging per order ($, before tax)', USD, 0.55, 'Wrapper, fry cup, bag and napkins of a take-out order.' + EST),
    ('cp', 'Marketplace packaging per order ($, before tax)', USD, 0.75, 'Vented fry box and tamper-evident seal.' + EST),
    ('rec', 'Price markup on the marketplace (%)', PCT, 0.15, 'What you raise the price on the marketplace to compensate.' + EST),
    ('com', 'Delivery marketplace commission (%)', PCT, 0.25, 'On the price before tax; check your contract.' + EST),
]
E = OrderedDict((k[0], 5 + i) for i, k in enumerate(ESC_IN))
E_BO, E_PA = 5 + len(ESC_IN), 5 + len(ESC_IN) + 1                                # enlaces a las hojas Beef & Patty y Fries & Oil
E_CAB = E_PA + 2
#: producto: (nombre, balls, lonchas de queso, buns, otros ingredientes $, porciones de fries, precio de carta, vendidos a la semana)
PRODS = [
    ('Single smash burger (1 ball)', 1, 1, 1, 0.30, 0, 8.49, 230),
    ('Double smash burger (2 balls)', 2, 2, 1, 0.35, 0, 10.99, 390),
    ('Triple smash burger (3 balls)', 3, 3, 1, 0.40, 0, 13.49, 65),
    ('Double bacon smash burger', 2, 2, 1, 0.80, 0, 12.49, 150),
    ('Oklahoma onion smash burger (2 balls)', 2, 2, 1, 0.45, 0, 11.99, 80),
    ('House smash burger (2 balls)', 2, 2, 1, 0.85, 0, 12.99, 115),
    ('Crispy chicken sandwich', 0, 1, 1, 2.10, 0, 10.49, 100),
    ('Veggie burger', 0, 1, 1, 1.90, 0, 9.99, 35),
    ('Fries', 0, 0, 0, 0.0, 1, 3.99, 280),
    ('Loaded fries (cheese sauce and bacon)', 0, 0, 0, 0.85, 1, 6.49, 150),
    ('Shake', 0, 0, 0, 1.45, 0, 6.49, 100),
    ('Dessert (cookie or sundae)', 0, 0, 0, 1.10, 0, 4.49, 90),
    ('Combo: double + fries + soft drink', 2, 2, 1, 0.60, 1, 15.99, 380),
    ('Meal for two: 2 doubles + 2 fries', 4, 4, 2, 0.70, 2, 28.99, 55),
]
NPR = len(PRODS)
J_SIM, J_DOB, J_COMBO = 0, 1, 12
E_P0 = E_CAB + 1
E_P1 = E_P0 + NPR - 1
E_T0 = E_P1 + 2
TOT = [
    ('ing', 'Weekly sales before tax ($)', USD, 'SUMPRODUCT(%s,%s)'),
    ('cos', 'Food cost of the week ($)', USD, 'SUMPRODUCT(%s,%s)'),
    ('fcp', 'WEIGHTED FOOD COST (%)', PCT, None),
    ('semp', 'Traffic light of the weighted food cost', 'General', None),
]
T = OrderedDict((k[0], E_T0 + 1 + i) for i, k in enumerate(TOT))
E_FIN = E_T0 + len(TOT) + 2
# columnas de la tabla
CN, CBO, CLQ, CPN, COT, CRA, CPVP, CUD = 'A', 'B', 'C', 'D', 'E', 'F', 'G', 'H'
CCO, CPS, CFC, CSE, CSG, CMS, CML, CMP, CRK = 'I', 'J', 'K', 'L', 'M', 'N', 'O', 'P', 'Q'
SALP = [('c_tot', CCO, USD3), ('pvp_sin', CPS, USD), ('fc', CFC, PCT), ('sem', CSE, 'General'), ('pvp_sug', CSG, USD),
        ('sala', CMS, USD), ('llevar', CML, USD), ('plat', CMP, USD), ('rank', CRK, ENT)]

EXPUESTOS = OrderedDict([
    ('COST_LB_USABLE_BLEND', (HC, 'B%d' % S_C['cku'])),
    ('COST_LB_USABLE_CHUCK', (HC, 'C%d' % S_C['cku'])),
    ('GRIND_PAYS_OFF_BELOW', (HC, 'B%d' % S_D['X'])),
    ('BALLS_PER_LB', (HC, 'B%d' % S_D['bkg'])),
    ('BALLS_PER_LB_NET', (HC, 'B%d' % S_D['bkn'])),
    ('BEEF_PER_DAY_LB', (HC, 'B%d' % S_D['carne_dia'])),
    ('COST_SINGLE', (HE, '%s%d' % (CCO, E_P0 + J_SIM))),
    ('COST_DOUBLE', (HE, '%s%d' % (CCO, E_P0 + J_DOB))),
    ('FC_SINGLE', (HE, '%s%d' % (CFC, E_P0 + J_SIM))),
    ('FC_DOUBLE', (HE, '%s%d' % (CFC, E_P0 + J_DOB))),
    ('FC_COMBO', (HE, '%s%d' % (CFC, E_P0 + J_COMBO))),
    ('FC_WEIGHTED', (HE, 'B%d' % T['fcp'])),
    ('COST_FROZEN_FRIES_PORTION', (HP, 'B%d' % S_A['c_cong'])),
    ('COST_FRESH_FRIES_PORTION', (HP, 'B%d' % S_A['c_fre'])),
    ('OIL_PER_PORTION', (HP, 'B%d' % S_A['ace_cong'])),
    ('OIL_LB_MONTH', (HP, 'B%d' % S_A['l_mes'])),
    ('MARGIN_DINEIN_COMBO', (HE, '%s%d' % (CMS, E_P0 + J_COMBO))),
    ('MARGIN_TAKEOUT_COMBO', (HE, '%s%d' % (CML, E_P0 + J_COMBO))),
    ('MARGIN_MARKETPLACE_COMBO', (HE, '%s%d' % (CMP, E_P0 + J_COMBO))),
])


def valores_defecto():
    """Todas las entradas por defecto en un solo dict plano (claves = las de las celdas verdes)."""
    v = OrderedDict()
    for i, (_n, pr, ut, tl, gr) in enumerate(CARNES):
        v['pr%d' % i], v['ut%d' % i], v['tl%d' % i], v['gr%d' % i] = pr, ut, tl, gr
    for ent in (ENT_C, ENT_A):
        for k, _r, _f2, val, _c, _n in ent:
            v[k] = val
    for k, _r, _f2, val, _n in ESC_IN:
        v[k] = val
    for j, (_nom, bo, lq, pn, ot, ra, pvp, ud) in enumerate(PRODS):
        v['bo%d' % j], v['lq%d' % j], v['pn%d' % j], v['ot%d' % j] = bo, lq, pn, ot
        v['ra%d' % j], v['pvp%d' % j], v['ud%d' % j] = ra, pvp, ud
    return v


# ---------------------------------------------------------------------------------------------------------------------
# Cálculo en Python (espejo de las fórmulas) — fuente del autotest
# ---------------------------------------------------------------------------------------------------------------------
def calcular(v):
    r = OrderedDict()
    cols = []
    for i in range(NPE):
        c = OrderedDict()
        c['cku'] = v['pr%d' % i] / v['ut%d' % i] / (1 - v['m'])
        c['cku_mo'] = (v['pr%d' % i] + v['tl%d' % i] / v['MH'] * v['hh']) / v['ut%d' % i] / (1 - v['m'])
        c['c_bola'] = c['cku_mo'] * v['gb'] / v['OZ']
        c['c_dos'] = c['c_bola'] + c['c_bola']
        c['c_tres'] = c['c_bola'] + c['c_bola'] + c['c_bola']
        cols.append(c)
    r['Pe'] = cols
    d = OrderedDict()
    d['bkg'] = v['OZ'] / v['gb']
    d['bkn'] = v['OZ'] * (1 - v['m']) / v['gb']
    d['gser'] = v['gb'] * (1 - v['mc'])
    d['bdia'] = v['cd'] * v['bc']
    d['carne_dia'] = d['bdia'] * v['gb'] / v['OZ']
    d['lab_kg'] = cols[1]['cku_mo'] - cols[1]['cku']
    d['X'] = cols[0]['cku_mo'] - cols[1]['cku']
    d['sem'] = SEM_PICAR if d['lab_kg'] < d['X'] else SEM_BLEND
    d['rem'] = OPC[[o for o, _t in OPC].index(v['cong'])][1]
    r['D'] = d
    a = OrderedDict()
    a['cd_c'] = v['cd']
    a['aceite_cambio'] = v['lc'] * v['nc'] * v['pa']
    a['ace_cong'] = a['aceite_cambio'] / v['rcm'] + v['gc'] * v['ac'] / v['OZ'] * v['pa']
    a['pat_cong'] = v['gc'] * v['pc'] / v['OZ']
    a['c_cong'] = a['pat_cong'] + a['ace_cong']
    a['kg_bruto'] = v['gf'] / v['uf'] / v['OZ']
    a['pat_fre'] = a['kg_bruto'] * v['pf']
    a['ace_fre'] = a['aceite_cambio'] / v['rcm'] + v['gf'] * v['af'] / v['OZ'] * v['pa']
    a['c_fre'] = a['pat_fre'] + a['ace_fre']
    a['mo_fre'] = v['tf'] / v['MH'] * v['hp'] * a['kg_bruto']
    a['c_fre_mo'] = a['c_fre'] + a['mo_fre']
    a['mo_kg'] = v['tf'] / v['MH'] * v['hp']
    a['X'] = (a['c_cong'] - a['c_fre']) / a['kg_bruto']
    a['sem'] = SEM_FRE if a['mo_kg'] < a['X'] else SEM_CONG
    a['rac_dia'] = a['cd_c'] * v['rpc']
    a['kg_cong_dia'] = a['rac_dia'] * v['gc'] / v['OZ']
    a['kg_fre_dia'] = a['rac_dia'] * a['kg_bruto']
    a['l_mes'] = a['rac_dia'] * v['dm'] * (v['lc'] * v['nc'] / v['rcm'] + v['gc'] * v['ac'] / v['OZ'])
    r['A'] = a
    sel, tp = int(v['sel']), int(v['tp'])
    r['bola'] = cols[sel - 1]['c_bola']
    r['pata'] = a['c_cong'] if tp == 1 else a['c_fre_mo']
    ps = []
    for j in range(NPR):
        q = OrderedDict()
        q['c_tot'] = (v['bo%d' % j] * r['bola'] + v['pn%d' % j] * v['pan'] + v['lq%d' % j] * v['qs']
                      + v['ot%d' % j] + v['ra%d' % j] * r['pata'])
        q['pvp_sin'] = v['pvp%d' % j] / (1 + v['tax'])
        q['fc'] = q['c_tot'] / q['pvp_sin']
        q['sem'] = FC_OK if q['fc'] <= v['obj'] else (FC_VIG if q['fc'] <= v['obj'] + v['tol'] else FC_ALTO)
        q['pvp_sug'] = q['c_tot'] / v['obj'] * (1 + v['tax'])
        q['sala'] = q['pvp_sin'] - q['c_tot']
        q['llevar'] = q['sala'] - v['cl']
        q['plat'] = q['pvp_sin'] * (1 + v['rec']) * (1 - v['com']) - q['c_tot'] - v['cp']
        ps.append(q)
    for q in ps:
        q['rank'] = 1 + sum(1 for o in ps if o['sala'] > q['sala'])
    r['Pr'] = ps
    t = OrderedDict()
    t['ing'] = sum(q['pvp_sin'] * v['ud%d' % j] for j, q in enumerate(ps))
    t['cos'] = sum(q['c_tot'] * v['ud%d' % j] for j, q in enumerate(ps))
    t['fcp'] = 0 if t['ing'] == 0 else t['cos'] / t['ing']
    t['semp'] = FC_OK if t['fcp'] <= v['obj'] else (FC_VIG if t['fcp'] <= v['obj'] + v['tol'] else FC_ALTO)
    r['T'] = t
    return r


# ---------------------------------------------------------------------------------------------------------------------
# Helpers (de la Taquería, la Pollería y el Sushi)
# ---------------------------------------------------------------------------------------------------------------------
def _f(ws, coord, v, b=False, sz=11, i=False, fmt=None, wrap=False, fill=None, locked=True):
    c = ws[coord]
    c.value = v
    c.font = Font(name=FUENTE, size=sz, bold=b, italic=i)
    if fmt:
        c.number_format = fmt
    c.alignment = Alignment(wrap_text=wrap, vertical='top')
    if fill:
        c.fill = fill
    c.protection = Protection(locked=locked)
    return c


def _nota(ws, fila, texto, desde='E', hasta='F', chars=80.0):
    if not texto:
        return
    _f(ws, '%s%d' % (desde, fila), texto, wrap=True)
    if desde != hasta:
        ws.merge_cells('%s%d:%s%d' % (desde, fila, hasta, fila))
    ws.row_dimensions[fila].height = max(ws.row_dimensions[fila].height or 0, 16, 15 * math.ceil(len(texto) / chars))


def _imprimir(ws, horizontal=False):
    ws.page_setup.paperSize = M.PAPEL_EN                                          # 1 = US Letter
    ws.page_setup.orientation = 'landscape' if horizontal else 'portrait'
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.oddFooter.center.text = FOOTER
    ws.print_options.horizontalCentered = True
    ws.protection.sheet = True                                                    # sin contraseña
    ws.sheet_view.showGridLines = False


def _dv(ws, tipo, celdas, f1=None, f2=None, op=None, titulo='Invalid value', msg=''):
    dv = DataValidation(type=tipo, formula1=f1, formula2=f2, operator=op, allow_blank=False, showErrorMessage=True,
                        errorTitle=titulo, error=msg)
    ws.add_data_validation(dv)
    for c in celdas:
        dv.add(c)


def a_excel(expr, ref):
    return re.sub(r'\{(\w+)\}', lambda m: ref[m.group(1)], expr)


def _bloque(ws, titulo_in, ent, m, ra, sal, s, neg, chars=105.0, extra=None):
    """Escribe el bloque de entradas (verdes/grises) y el de resultados de una hoja de dos bloques."""
    _f(ws, 'A5', titulo_in, b=True, fill=GRIS)
    _f(ws, 'B5', 'Value', b=True, fill=GRIS)
    _f(ws, 'C5', 'Notes', b=True, fill=GRIS)
    for k, rot, fmt, val, const, nota in ent:
        f = m[k]
        _f(ws, 'A%d' % f, rot)
        if const:
            _f(ws, 'B%d' % f, val, fmt=fmt, fill=GRIS)
        else:
            _f(ws, 'B%d' % f, val, fmt=fmt, fill=VERDE, locked=False)
        _nota(ws, f, nota, desde='C', hasta='I', chars=chars)
    ref = dict((k, '$B$%d' % f) for k, f in m.items())
    ref.update((k, '$B$%d' % f) for k, f in s.items())
    ref.update(extra or {})
    for col, t in zip('ABC', ('RESULTS', 'Value', 'Notes')):
        _f(ws, '%s%d' % (col, ra), t, b=True, fill=GRIS)
    for k, rot, fmt, expr, nota in sal:
        f = s[k]
        _f(ws, 'A%d' % f, rot, b=k in neg, wrap=len(rot) > 60)
        _f(ws, 'B%d' % f, '=' + a_excel(expr, ref), fmt=fmt, b=k in neg, wrap=(fmt == 'General'))
        _nota(ws, f, nota, desde='C', hasta='I', chars=chars)


def _semaforo(ws, celda, ok, ko):
    m = re.match(r'([A-Z]+)(\d+)', celda)
    ws.conditional_formatting.add(celda, FormulaRule(formula=['$%s$%s="%s"' % (m.group(1), m.group(2), ok)], fill=OK))
    ws.conditional_formatting.add(celda, FormulaRule(formula=['$%s$%s="%s"' % (m.group(1), m.group(2), ko)], fill=ROJO))


# ---------------------------------------------------------------------------------------------------------------------
# Hojas
# ---------------------------------------------------------------------------------------------------------------------
def hoja_carne(wb):
    ws = wb.active
    ws.title = HC
    for col, w in (('A', 62), ('B', 19), ('C', 19), ('D', 19), ('E', 19), ('F', 72)):
        ws.column_dimensions[col].width = w
    _f(ws, 'A1', 'BEEF & PATTY CALCULATOR — COST PER USABLE LB, PER BALL, BALLS PER LB AND LB PER DAY', b=True, sz=14)
    _f(ws, 'A2', 'Type only in the GREEN cells: everything else is calculated. Amounts are before sales tax. Every default is a Kit estimate, '
                 "editable: replace it with your distributor's price and, above all, with what you measure on your own balls.", wrap=True)
    ws.merge_cells('A2:F2')
    ws.row_dimensions[2].height = 32
    _f(ws, 'A3', 'FIXED NOTICE. Ground beef (FDA Food Code 2022): keep raw ground beef at 41 °F (5 °C) or below and store it below ready-to-eat '
                 'food (3-501.16 and 3-302.11). A house-grind log is good practice and useful in a recall; ask your health department what '
                 'it requires. The minutes to trim and grind (3.0) and the cook loss have no source: weigh and time them.',
       wrap=True, b=True, fill=GRIS)
    ws.merge_cells('A3:F3')
    ws.row_dimensions[3].height = 64
    # A. tres carnes y una libre
    _f(ws, 'A5', 'A. THREE BEEFS AND ONE FREE COLUMN', b=True, fill=GRIS)
    for i, c in enumerate(CARNES):
        _f(ws, '%s5' % PC[i], c[0], b=True, fill=VERDE, wrap=True, locked=False)
    _f(ws, 'F5', 'Notes', b=True, fill=GRIS)
    ws.row_dimensions[5].height = 62
    for k, rot, fmt, idx in (('pr', 'Purchase price ($/lb, before tax)', USD, 1), ('ut', 'Usable share after trimming and grinding (% of the weight you buy)', PCT, 2),
                             ('tl', 'Minutes to trim and grind 1 lb you buy (min)', DEC, 3),
                             ('gr', 'Fat of the beef (%, informational)', PCT, 4)):
        f = P_ROW[k]
        _f(ws, 'A%d' % f, rot, wrap=True)
        for i, p in enumerate(CARNES):
            _f(ws, '%s%d' % (PC[i], f), p[idx], fmt=fmt, fill=VERDE, locked=False)
        _nota(ws, f, P_NOTAS[k], desde='F', hasta='F', chars=75.0)
    # B. datos comunes
    _f(ws, 'A11', 'B. COMMON INPUTS', b=True, fill=GRIS)
    _f(ws, 'B11', 'Value', b=True, fill=GRIS)
    _f(ws, 'C11', 'Notes', b=True, fill=GRIS)
    for k, rot, fmt, val, const, nota in ENT_C:
        f = M_C[k]
        _f(ws, 'A%d' % f, rot)
        if const:
            _f(ws, 'B%d' % f, val, fmt=fmt, fill=GRIS)
        else:
            _f(ws, 'B%d' % f, val, fmt=fmt, fill=VERDE, locked=False, wrap=(k == 'cong'))
        _nota(ws, f, nota, desde='C', hasta='F', chars=125.0)
    ws.merge_cells('B%d:E%d' % (M_C['cong'], M_C['cong']))
    ws.row_dimensions[M_C['cong']].height = 32
    # C. coste por lb útil
    _f(ws, 'A%d' % PE_RC, 'C. COST PER USABLE LB AND PER BALL', b=True, fill=GRIS)
    for i in range(NPE):
        _f(ws, '%s%d' % (PC[i], PE_RC), '=%s5' % PC[i], b=True, fill=GRIS, wrap=True)
    _f(ws, 'F%d' % PE_RC, 'Notes', b=True, fill=GRIS)
    ws.row_dimensions[PE_RC].height = 62
    for k, rot, fmt, expr, nota in SAL_C:
        f = S_C[k]
        _f(ws, 'A%d' % f, rot, b=k in ('cku', 'cku_mo'))
        for i in range(NPE):
            c = PC[i]
            ref = dict((kk, '$B$%d' % ff) for kk, ff in M_C.items())
            ref.update((kk, '%s$%d' % (c, ff)) for kk, ff in P_ROW.items())
            ref.update((kk, '%s$%d' % (c, ff)) for kk, ff in S_C.items())
            _f(ws, '%s%d' % (c, f), '=' + a_excel(expr, ref), fmt=fmt, b=k in ('cku', 'cku_mo'))
        _nota(ws, f, nota, desde='F', hasta='F', chars=75.0)
    # D. balls, lb al día y grind or buy
    _f(ws, 'A%d' % PE_RD, 'D. BALLS PER LB, LB PER DAY AND GRIND IN-HOUSE?', b=True, fill=GRIS)
    _f(ws, 'B%d' % PE_RD, 'Value', b=True, fill=GRIS)
    _f(ws, 'C%d' % PE_RD, 'Notes', b=True, fill=GRIS)
    ref = dict((k, '$B$%d' % f) for k, f in M_C.items())
    ref.update((k, '$B$%d' % f) for k, f in S_D.items())
    ref.update({'cku0': '$B$%d' % S_C['cku'], 'cku1': '$C$%d' % S_C['cku'], 'cku_mo0': '$B$%d' % S_C['cku_mo'],
                'cku_mo1': '$C$%d' % S_C['cku_mo'],
                'rem_r': '$B$%d:$B$%d' % (PE_O0, PE_O1), 'opc_r': '$A$%d:$A$%d' % (PE_O0, PE_O1)})
    neg = ('bkg', 'carne_dia', 'X', 'sem')
    for k, rot, fmt, expr, nota in SAL_D:
        f = S_D[k]
        _f(ws, 'A%d' % f, rot, b=k in neg, wrap=True)
        _f(ws, 'B%d' % f, '=' + a_excel(expr, ref), fmt=fmt, b=k in neg, wrap=(fmt == 'General'))
        _nota(ws, f, nota, desde='C', hasta='F', chars=125.0)
    ws.merge_cells('B%d:E%d' % (S_D['rem'], S_D['rem']))
    ws.row_dimensions[S_D['rem']].height = 48
    ws.row_dimensions[S_D['X']].height = 32
    ws.row_dimensions[S_D['sem']].height = 32
    _semaforo(ws, 'B%d' % S_D['sem'], SEM_PICAR, SEM_BLEND)
    _f(ws, 'A%d' % PE_NOTA, NOTA_COCCION, b=True, wrap=True)
    ws.merge_cells('A%d:F%d' % (PE_NOTA, PE_NOTA))
    ws.row_dimensions[PE_NOTA].height = 32
    # lista del desplegable
    _f(ws, 'A%d' % PE_RO, 'GROUND BEEF DROPDOWN LIST', b=True, fill=GRIS)
    _f(ws, 'B%d' % PE_RO, 'Reminder', b=True, fill=GRIS)
    for i, (op, rem) in enumerate(OPC):
        f = PE_O0 + i
        _f(ws, 'A%d' % f, op, wrap=True)
        _f(ws, 'B%d' % f, rem, wrap=True)
        ws.merge_cells('B%d:F%d' % (f, f))
        ws.row_dimensions[f].height = 48
    # validaciones
    _dv(ws, 'decimal', ['%s%d' % (c, P_ROW['pr']) for c in PC], f1='0', op='greaterThan', msg='Type a number greater than 0.')
    _dv(ws, 'decimal', ['%s%d' % (c, P_ROW['ut']) for c in PC], f1='0.3', f2='1', op='between', titulo='Invalid percentage',
        msg='Type a value between 0.3 and 1 (0.9 = 90%).')
    _dv(ws, 'decimal', ['%s%d' % (c, P_ROW['tl']) for c in PC], f1='0', op='greaterThanOrEqual', msg='Type a number greater than or equal to 0.')
    _dv(ws, 'decimal', ['%s%d' % (c, P_ROW['gr']) for c in PC] + ['B%d' % M_C['m'], 'B%d' % M_C['mc']], f1='0', f2='0.5', op='between',
        titulo='Invalid percentage', msg='Type a percentage between 0 and 0.5 (0.05 = 5%).')
    _dv(ws, 'decimal', ['B%d' % M_C[k] for k in ('gb', 'bc')], f1='0', op='greaterThan', msg='Type a number greater than 0.')
    _dv(ws, 'decimal', ['B%d' % M_C['hh']], f1='0', op='greaterThanOrEqual', msg='Type a number greater than or equal to 0.')
    _dv(ws, 'whole', ['B%d' % M_C['cd']], f1='1', f2='1000000', op='between', titulo='Invalid number', msg='Type a whole number of 1 or more.')
    _dv(ws, 'list', ['B%d' % M_C['cong']], f1='$A$%d:$A$%d' % (PE_O0, PE_O1), titulo='Invalid option', msg='Pick an option from the list.')
    _imprimir(ws, horizontal=True)


def hoja_patatas(wb):
    ws = wb.create_sheet(HP)
    for col, w in (('A', 66), ('B', 15), ('C', 15), ('D', 16), ('E', 14), ('F', 14), ('G', 14), ('H', 12), ('I', 15)):
        ws.column_dimensions[col].width = w
    _f(ws, 'A1', 'FRIES & OIL CALCULATOR — COST PER PORTION, LB PER DAY AND LB OF OIL PER MONTH', b=True, sz=14)
    _f(ws, 'A2', 'Type only in the GREEN cells: everything else is calculated. Amounts are before sales tax. Every default is a Kit estimate, '
                 'editable: weigh your portions, time the peeling and replace the figures with yours.', wrap=True)
    ws.merge_cells('A2:I2')
    ws.row_dimensions[2].height = 32
    _f(ws, 'A3', 'FIXED NOTICE. Frying oil: there is no federal limit on total polar materials (TPM); set the change by your own readings and '
                 'the manufacturer\'s guidance. A shared fryer is an allergen risk: keep the breaded chicken out of the fries fryers.',
       wrap=True, b=True, fill=GRIS)
    ws.merge_cells('A3:I3')
    ws.row_dimensions[3].height = 36
    _bloque(ws, 'A. THE POTATOES AND THE OIL', ENT_A, M_A, M_RA, SAL_A, S_A, ('c_cong', 'c_fre', 'X', 'sem', 'rac_dia', 'l_mes', 'ace_cong'),
            extra={'cd_carne': "'%s'!$B$%d" % (HC, M_C['cd'])})
    ws.row_dimensions[S_A['X']].height = 32
    ws.row_dimensions[S_A['sem']].height = 32
    _semaforo(ws, 'B%d' % S_A['sem'], SEM_FRE, SEM_CONG)
    _f(ws, 'A%d' % PA_NOTA, NOTA_FRITURA, b=True, wrap=True)
    ws.merge_cells('A%d:I%d' % (PA_NOTA, PA_NOTA))
    ws.row_dimensions[PA_NOTA].height = 48
    cel = lambda k: 'B%d' % M_A[k]
    _dv(ws, 'decimal', [cel(k) for k in ('pc', 'gc', 'pf', 'gf', 'hp', 'lc', 'pa', 'rcm', 'rpc')], f1='0', op='greaterThan',
        msg='Type a number greater than 0.')
    _dv(ws, 'decimal', [cel('uf')], f1='0.3', f2='1', op='between', titulo='Invalid percentage', msg='Type a value between 0.3 and 1 (0.75 = 75%).')
    _dv(ws, 'decimal', [cel('tf')], f1='0', op='greaterThanOrEqual', msg='Type a number greater than or equal to 0.')
    _dv(ws, 'decimal', [cel('ac'), cel('af')], f1='0', f2='0.5', op='between', titulo='Invalid percentage',
        msg='Type a percentage between 0 and 0.5 (0.05 = 5%).')
    _dv(ws, 'whole', [cel('nc')], f1='1', f2='20', op='between', titulo='Invalid number', msg='Type a whole number from 1 to 20.')
    _dv(ws, 'whole', [cel('dm')], f1='1', f2='31', op='between', titulo='Invalid number', msg='Type a whole number from 1 to 31.')
    _imprimir(ws, horizontal=True)


def hoja_escandallo(wb):
    ws = wb.create_sheet(HE)
    for col, w in ((CN, 44), (CBO, 11), (CLQ, 11), (CPN, 10), (COT, 13), (CRA, 12), (CPVP, 13), (CUD, 14), (CCO, 12), (CPS, 12), (CFC, 11),
                   (CSE, 17), (CSG, 14), (CMS, 12), (CML, 12), (CMP, 13), (CRK, 10)):
        ws.column_dimensions[col].width = w
    _f(ws, 'A1', 'BURGER RESTAURANT COSTING — 14 PRODUCTS, PER PRODUCT AND PER CHANNEL', b=True, sz=14)
    _f(ws, 'A2', 'Type only in the GREEN cells (names, balls, cheese, buns, other ingredients, fries, menu prices and units). Amounts are before '
                 'sales tax except the menu price. Every default is a Kit estimate, editable: type yours. The cost of the ball comes from '
                 'the "%s" sheet, the cost of a portion of fries from the "%s" sheet.' % (HC, HP), wrap=True)
    ws.merge_cells('A2:%s2' % CRK)
    ws.row_dimensions[2].height = 32
    for col, t in zip('ABD', ('GENERAL INPUTS', 'Value', 'Notes')):
        _f(ws, '%s4' % col, t, b=True, fill=GRIS)
    for k, rot, fmt, val, nota in ESC_IN:
        f = E[k]
        _f(ws, 'A%d' % f, rot)
        _f(ws, 'B%d' % f, val, fmt=fmt, fill=VERDE, locked=False)
        _nota(ws, f, nota, desde='D', hasta=CRK, chars=150.0)
    _f(ws, 'A%d' % E_BO, 'Cost of the beef of one ball ($, from the Beef & Patty sheet)')
    _f(ws, 'B%d' % E_BO, "=CHOOSE($B$%d,%s)" % (E['sel'], ','.join("'%s'!$%s$%d" % (HC, c, S_C['c_bola']) for c in PC)), fmt=USD3)
    _f(ws, 'A%d' % E_PA, 'Cost of one portion of fries ($, from the Fries & Oil sheet)')
    _f(ws, 'B%d' % E_PA, "=CHOOSE($B$%d,'%s'!$B$%d,'%s'!$B$%d)" % (E['tp'], HP, S_A['c_cong'], HP, S_A['c_fre_mo']), fmt=USD3)
    cab = [(CN, 'PRODUCT'), (CBO, 'Balls (no.)'), (CLQ, 'Cheese slices (no.)'), (CPN, 'Buns (no.)'), (COT, 'Other ingredients ($)'),
           (CRA, 'Fries portions (no.)'), (CPVP, 'Menu price ($)'), (CUD, 'Sold per week (no.)'), (CCO, 'COST ($)'),
           (CPS, 'Price before tax ($)'), (CFC, 'FOOD COST (%)'), (CSE, 'Traffic light'), (CSG, 'Suggested menu price ($)'),
           (CMS, 'Margin dine-in ($)'), (CML, 'Margin take-out / pickup ($)'), (CMP, 'Margin marketplace ($)'), (CRK, 'Rank by margin')]
    for col, t in cab:
        _f(ws, '%s%d' % (col, E_CAB), t, b=True, fill=GRIS, wrap=True)
    ws.row_dimensions[E_CAB].height = 48
    for j, (nom, bo, lq, pn, ot, ra, pvp, ud) in enumerate(PRODS):
        f = E_P0 + j
        _f(ws, 'A%d' % f, nom, fill=VERDE, locked=False, b=True)
        for col, val, fmt in ((CBO, bo, '#,##0.0;-#,##0.0;"-"'), (CLQ, lq, '#,##0.0;-#,##0.0;"-"'), (CPN, pn, '#,##0.0;-#,##0.0;"-"'),
                              (COT, ot, USD), (CRA, ra, '#,##0.0;-#,##0.0;"-"'), (CUD, ud, ENT)):
            _f(ws, '%s%d' % (col, f), float(val), fmt=fmt, fill=VERDE, locked=False)
        _f(ws, '%s%d' % (CPVP, f), float(pvp), fmt=USD, fill=VERDE, locked=False, b=True)
        ex = {
            CCO: '%s%d*$B$%d+%s%d*$B$%d+%s%d*$B$%d+%s%d+%s%d*$B$%d' % (
                CBO, f, E_BO, CPN, f, E['pan'], CLQ, f, E['qs'], COT, f, CRA, f, E_PA),
            CPS: '%s%d/(1+$B$%d)' % (CPVP, f, E['tax']),
            CFC: '%s%d/%s%d' % (CCO, f, CPS, f),
            CSE: 'IF(%s%d<=$B$%d,"%s",IF(%s%d<=$B$%d+$B$%d,"%s","%s"))' % (
                CFC, f, E['obj'], FC_OK, CFC, f, E['obj'], E['tol'], FC_VIG, FC_ALTO),
            CSG: '%s%d/$B$%d*(1+$B$%d)' % (CCO, f, E['obj'], E['tax']),
            CMS: '%s%d-%s%d' % (CPS, f, CCO, f),
            CML: '%s%d-$B$%d' % (CMS, f, E['cl']),
            CMP: '%s%d*(1+$B$%d)*(1-$B$%d)-%s%d-$B$%d' % (CPS, f, E['rec'], E['com'], CCO, f, E['cp']),
            CRK: '1+COUNTIF($%s$%d:$%s$%d,">"&%s%d)' % (CMS, E_P0, CMS, E_P1, CMS, f),
        }
        for _k, col, fmt in SALP:
            neg = col in (CCO, CFC)
            _f(ws, '%s%d' % (col, f), '=' + ex[col], fmt=fmt, b=neg or col == CRK, wrap=(col == CSE))
        ws.row_dimensions[f].height = 32
    rng = '%s%d:%s%d' % (CSE, E_P0, CSE, E_P1)
    top = '%s%d' % (CSE, E_P0)
    for txt, fill in ((FC_OK, OK), (FC_VIG, AMBAR), (FC_ALTO, ROJO)):
        ws.conditional_formatting.add(rng, FormulaRule(formula=['%s="%s"' % (top, txt)], fill=fill))
    # ponderado
    _f(ws, 'A%d' % E_T0, 'WEIGHTED FOOD COST OF THE WEEK (sum of costs × units ÷ sum of prices before tax × units)', b=True, fill=GRIS)
    _f(ws, 'B%d' % E_T0, 'Value', b=True, fill=GRIS)
    rc = '$%s$%d:$%s$%d' % (CCO, E_P0, CCO, E_P1)
    rp = '$%s$%d:$%s$%d' % (CPS, E_P0, CPS, E_P1)
    ru = '$%s$%d:$%s$%d' % (CUD, E_P0, CUD, E_P1)
    exs = {'ing': 'SUMPRODUCT(%s,%s)' % (rp, ru), 'cos': 'SUMPRODUCT(%s,%s)' % (rc, ru),
           'fcp': 'IF($B$%d=0,0,$B$%d/$B$%d)' % (T['ing'], T['cos'], T['ing']),
           'semp': 'IF($B$%d<=$B$%d,"%s",IF($B$%d<=$B$%d+$B$%d,"%s","%s"))' % (
               T['fcp'], E['obj'], FC_OK, T['fcp'], E['obj'], E['tol'], FC_VIG, FC_ALTO)}
    for k, rot, fmt, _x in TOT:
        f = T[k]
        _f(ws, 'A%d' % f, rot, b=k == 'fcp')
        _f(ws, 'B%d' % f, '=' + exs[k], fmt=fmt, b=k == 'fcp', wrap=(k == 'semp'))
    ws.row_dimensions[T['semp']].height = 32
    for txt, fill in ((FC_OK, OK), (FC_VIG, AMBAR), (FC_ALTO, ROJO)):
        ws.conditional_formatting.add('B%d' % T['semp'], FormulaRule(formula=['$B$%d="%s"' % (T['semp'], txt)], fill=fill))
    _f(ws, 'A%d' % E_FIN, 'The food cost in the table is the product served dine-in. On a take-out order the packaging is deducted; on the '
                          'marketplace the price is raised by the markup and the commission is deducted on that price, plus the packaging. '
                          'Margins are ranked by the dine-in margin. Other ingredients: toppings and sauce (pickles, onion, lettuce, special '
                          'sauce) about $0.35 on a double; bacon, the breaded chicken breast, the veggie patty, the shake base and the soft '
                          'drink (syrup and cup) at Kit estimate prices. FIXED NOTICE: a smash burger travels worse than other food: the '
                          'crust softens in about 15 minutes. Use vented boxes for the fries and price the marketplace menu about 15% higher. '
                          'This cost is THEORETICAL: the plan\'s 34% food cost ("0. Assumptions" B13 of the Burger Restaurant Financial '
                          'Projections) is the real usage of the whole food menu, with the balls, buns and sauces thrown out, staff meals, '
                          'comps, mistakes and marketplace refunds and promotions: this weighted theoretical cost comes out lower. Do not '
                          'copy it into the plan.', i=True, wrap=True)
    ws.merge_cells('A%d:%s%d' % (E_FIN, CRK, E_FIN))
    ws.row_dimensions[E_FIN].height = 96
    ws.freeze_panes = 'B%d' % (E_CAB + 1)
    # validaciones
    _dv(ws, 'decimal', ['B%d' % E[k] for k in ('tax', 'tol', 'rec', 'com')], f1='0', f2='0.9', op='between',
        titulo='Invalid percentage', msg='Type a percentage between 0 and 0.9 (0.1 = 10%).')
    _dv(ws, 'decimal', ['B%d' % E['obj']], f1='0.01', f2='0.9', op='between', titulo='Invalid target',
        msg='The target cannot be 0 (it is a divisor): type a value between 0.01 and 0.9.')
    _dv(ws, 'decimal', ['B%d' % E[k] for k in ('pan', 'qs', 'cl', 'cp')]
        + ['%s%d' % (c, f) for c in (CBO, CLQ, CPN, COT, CRA, CPVP, CUD) for f in range(E_P0, E_P1 + 1)],
        f1='0', op='greaterThanOrEqual', titulo='Invalid amount', msg='Type a number greater than or equal to 0.')
    _dv(ws, 'whole', ['B%d' % E['sel']], f1='1', f2=str(NPE), op='between', titulo='Invalid beef',
        msg='Type 1, 2, 3 or 4, the number of the column on the Beef & Patty sheet.')
    _dv(ws, 'whole', ['B%d' % E['tp']], f1='1', f2='2', op='between', titulo='Invalid fries', msg='Type 1 (frozen) or 2 (fresh-cut).')
    _imprimir(ws, horizontal=True)


def _n(x, d=2):
    """Número en formato EN (punto decimal, coma de millares)."""
    return ('{:,.%df}' % d).format(x)


def instrucciones():
    """Texto de la hoja Instructions; las cifras salen de calcular() con los valores por defecto (no se desfasan)."""
    r = calcular(valores_defecto())
    pe, d, a, pr, t = r['Pe'], r['D'], r['A'], r['Pr'], r['T']
    combo = pr[J_COMBO]
    return [
        ('INSTRUCTIONS — Smash Burger Patty, Fries & Costing Calculator', True),
        ('Three tools in one workbook: the "Beef & Patty" sheet tells you what a usable lb of ground chuck blend, whole chuck and pre-portioned balls '
         'costs, what each ball is worth, how many balls make a lb, how many lb of beef you need per day and whether grinding in-house pays off; '
         'the "Fries & Oil" sheet calculates the cost of a frozen and a fresh-cut portion, the oil per portion, the lb of potatoes per day and the '
         'lb of oil per month; the "Costing" sheet calculates the cost, food cost, margin and menu price of 14 products, with the margin dine-in, '
         'take-out and on the marketplace, and the weighted food cost of your week. You type only in the GREEN cells; everything else is '
         'calculated.', False),
        ("Every default is a Kit estimate, editable: coherent examples for the case of the business plan, not anyone's prices. Replace them with "
         "your distributor's prices and with what you measure on your flat-top and your fryers before you decide menu prices.", False),
        ('HOW THIS FILE FITS THE KIT', True),
        ('This calculator is one of four files: Burger Restaurant Business Plan, Burger Restaurant Financial Projections, Burger Restaurant Startup '
         'Checklist and this one. The cost per burger you get here is NOT copied into the plan: the plan uses the real usage of the whole menu '
         '(see "From theoretical food cost to the plan\'s 34%" below).', False),
        ('SHEET "BEEF & PATTY"', True),
        ('1. Weigh 10 balls and type the average in "Ounces per ball". Weigh also 10 balls before and after the flat-top to measure the cook loss '
         '(no source: measure it). If you grind in-house, weigh the whole chuck before trimming and the clean beef after, and time how long it '
         'takes you to trim and grind 3 lb.', False),
        ('2. The cost per usable lb is the price ÷ usable ÷ (1 − end-of-day shrink). With the Kit estimates it comes out at $%s per lb for the '
         'blend bought ground and $%s per lb for whole chuck, before labor. The cost of a ball is that cost per usable lb with labor × ounces ÷ 16: '
         '$%s with the blend.' % (_n(pe[0]['cku']), _n(pe[1]['cku']), _n(pe[0]['c_bola'], 3)), False),
        ('3. Balls per lb. Raw, it is 16 ÷ ounces per ball: %s balls of 3 oz per lb, the figure on the supplier spec sheet. Counting the 3%% '
         'end-of-day shrink you get %s. With 1.6 balls per cover and %s covers per day, you need about %s lb of beef per day: it is the daily '
         'order and, if you grind in-house, the day\'s grinding.' % (_n(d['bkg'], 1), _n(d['bkn'], 1), _n(195, 0), _n(d['carne_dia'], 1)), False),
        ('4. Grind in-house or buy ground? With the Kit estimates, grinding pays off only as long as trimming and grinding 1 lb cost less than $%s '
         'per usable lb; with 3 minutes per lb and $18.70 an hour it costs $%s: the verdict is "%s". Time your minutes before you buy a grinder.'
         % (_n(d['X']), _n(d['lab_kg']), d['sem']), False),
        ('5. Ground beef safety (FDA Food Code 2022). Cook every smash burger to 155 °F (68 °C) for 17 seconds (3-401.11(A)(2)); the other '
         'combinations in the Food Code are 150 °F (66 °C) for 1 minute, 145 °F (63 °C) for 3 minutes and 158 °F (70 °C) instantaneous. Probe one '
         'burger per batch: a thin smash burger cooks fast, but the crust is not a thermometer. Keep raw ground beef at 41 °F (5 °C) or below '
         '(3-501.16) and below ready-to-eat food (3-302.11). The consumer advisory (3-603.11: a disclosure and a reminder on the menu, the website '
         'and the marketplace listings) is needed only if you offer burgers cooked to order at the guest\'s request (3-401.11(D)), and never on the '
         "children's menu. The example serves every burger cooked through.", False),
        ('6. House-grind log. If you grind in-house, keep a log as good practice and for recalls: supplier establishment numbers and lots, and the '
         'date and time of every grind and of every cleaning and sanitizing of the grinder. Whether a log is required depends on your health '
         'department and on the FSIS rule for retail grinders (9 CFR 320.1(b)(4)): ask them. The dropdown on the sheet is a reminder only: it '
         'changes no figure.', False),
        ('SHEET "FRIES & OIL"', True),
        ('1. Weigh 10 portions of fries and use the average in ounces. If you cut fresh, weigh 10 lb unpeeled and what is left cut to measure the '
         'usable share, and time the peeling, cutting and soaking of 3 lb. For the oil, note how many portions each change lasts by your TPM '
         'readings (no source: measure it).', False),
        ('2. With the Kit estimates, a frozen portion costs $%s (of which $%s is oil) and a fresh-cut portion $%s without labor and $%s with it. '
         'Fresh-cut pays off only as long as peeling and cutting 1 lb cost less than $%s; with 3 minutes per lb and $18.70 an hour it costs '
         '$%s: the verdict is "%s".' % (_n(a['c_cong']), _n(a['ace_cong']), _n(a['c_fre']), _n(a['c_fre_mo']), _n(a['X']), _n(a['mo_kg']), a['sem']),
         False),
        ('3. With %s portions per day (covers × 0.8) you need about %s lb of frozen fries per day, or %s lb of fresh potatoes unpeeled, and about '
         '%s lb of oil per month.' % (_n(a['rac_dia'], 0), _n(a['kg_cong_dia'], 1), _n(a['kg_fre_dia'], 1), _n(a['l_mes'], 0)), False),
        ('4. Frying oil. There is no federal limit on total polar materials (TPM): set the change by your own readings and the manufacturer\'s '
         'guidance, filter the oil every night and send used oil to a licensed renderer. Fry at 350-355 °F. Acrylamide has no federal limit for '
         'restaurants; in California, ask your attorney about the Prop 65 warning for fried potatoes.', False),
        ('SHEET "COSTING"', True),
        ('1. Type the sales tax included in your menu prices (leave 0% if your menu is before tax), the food cost target, the watch band, which '
         'beef and which fries you use (the numbers of the previous sheets), the bun, the cheese per slice, the to-go and marketplace packaging, '
         'the marketplace markup and the commission.', False),
        ('2. In the table each row is a product. Type the balls, the slices of American cheese (one per ball), the buns, what the other '
         'ingredients cost (sauce, pickles, onion, butter for the bun and paper), the portions of fries, the menu price and the products you '
         'sell per week.', False),
        ('3. For each product you see the cost, the price before tax, the food cost with a traffic light, the menu price you would need to reach '
         'your target and the margin dine-in, take-out / pickup and on the marketplace, with the rank of each product by dine-in margin. Below, the '
         'weighted food cost of your week. With the Kit estimates, the double smash burger costs $%s and has a food cost of %s%%; the combo leaves '
         '$%s dine-in, $%s take-out and $%s on the marketplace. A smash burger travels worse than other food: the crust softens in about 15 '
         'minutes, so use vented boxes for the fries and price the marketplace menu about 15%% higher.'
         % (_n(pr[J_DOB]['c_tot']), _n(pr[J_DOB]['fc'] * 100, 1), _n(combo['sala']), _n(combo['llevar']), _n(combo['plat'])), False),
        ('ALLERGENS', True),
        ('The 9 major allergens in a burger: the bun (wheat; milk, egg, soy and sesame depending on the recipe), American cheese (milk), mayo '
         '(egg), shakes (milk), the breaded chicken (wheat, egg, milk) and the shared fryer oil. Mustard is not a major allergen in the US. Never '
         'promise "gluten-free": a shared fryer and a bun with wheat make it impossible to guarantee.', False),
        ("FROM THEORETICAL FOOD COST TO THE PLAN'S 34%", True),
        ("The food cost on this sheet is theoretical: what you would get if everything you buy were sold exactly as in the costing. With the Kit "
         'estimates the weighted week comes out around %s%%. The plan works with 34%% (sheet "0. Assumptions", cell B13, in the Burger Restaurant '
         'Financial Projections), the real usage of the whole food menu: between the recipe and the register you lose the balls, buns and sauces '
         'thrown out at close, staff meals, comps, mistakes and marketplace refunds and promotions. Your real comparison is this: the month\'s '
         'food purchases ÷ the month\'s food sales before tax. If after the first quarter it comes out above or below 34%%, update the food cost '
         'assumption in the Financial Projections. Do not copy the cost per burger of this calculator into the plan.' % _n(t['fcp'] * 100, 0),
         False),
        ('PRICES, SALES TAX AND SOURCES', True),
        ('US menus usually exclude sales tax, so the calculator works before tax and the sales tax input is 0%. Type your rate only if your menu '
         'prices include it. The prices and yields come from the kit\'s market research (October 2026): ground beef from BLS/FRED and retail '
         'reports, buns, cheese and fries from a restaurant supply store, the rest at Kit estimate prices. None of the figures in this calculator '
         'is a published figure for your business: replace them with yours.', False),
    ]


def hoja_instrucciones(wb):
    ws = wb.create_sheet(HI)
    ws.column_dimensions['A'].width = 110
    f = 1
    for t, b in instrucciones():
        _f(ws, 'A%d' % f, t, b=b, sz=13 if f == 1 else 11, wrap=True)
        if not b:
            ws.row_dimensions[f].height = max(16, 15 * math.ceil(len(t) / 120.0))
        f += 2 if b and f > 1 else 1
    f += 1
    for t in ('To edit a cell that is not green: Review → Unprotect Sheet (there is no password).',
              'More templates and kits in the catalog at ' + M.URL_TIENDA, M.FIRMA_EN, VERSION_LINEA, M.AVISO_EN):
        _f(ws, 'A%d' % f, t, wrap=True, i=t == M.AVISO_EN)
        ws.row_dimensions[f].height = max(16, 15 * math.ceil(len(t) / 120.0))
        f += 1
    _imprimir(ws)


def docprops(wb):
    pr = wb.properties
    pr.title = '%s · %s' % (TITULO, PRODUCTO)
    pr.subject = '%s · v%s' % (PRODUCTO, VERSION)
    pr.keywords = KEYWORDS
    pr.description = URL
    pr.category = 'AI Chef Pro · Digital products'
    pr.creator = pr.lastModifiedBy = 'AI Chef Pro'


def generar(destino):
    wb = openpyxl.Workbook()
    hoja_carne(wb)
    hoja_patatas(wb)
    hoja_escandallo(wb)
    hoja_instrucciones(wb)
    docprops(wb)
    os.makedirs(os.path.dirname(destino), exist_ok=True)
    wb.save(destino)
    import inject_cache as IC                                                      # noqa: E402
    nf, inj, fail = IC.inject(destino)
    print('%s: formulas=%d cache_inyectado=%d fallos_pycel=%d' % (os.path.basename(destino), nf, inj, fail))
    return fail


def valores_expuestos(ruta):
    """Lee de la CACHÉ del xlsx los valores expuestos: {nombre: valor}."""
    wv = openpyxl.load_workbook(ruta, data_only=True)
    return OrderedDict((n, wv[h][c].value) for n, (h, c) in EXPUESTOS.items())


# ---------------------------------------------------------------------------------------------------------------------
# Autotest
# ---------------------------------------------------------------------------------------------------------------------
#: restos del Sushi / Taco / Rotisserie / Pizza EN, euros, español, reglamentación ES y promesas (se suma datos_bur.RX_BUR_NO)
PROHIBIDOS = [r'sushi', r'sashimi', r'nigiri', r'\bnori\b', r'wasabi', r'tempura', r'gyoza', r'anisakis', r'surimi', r'sushiman',
              r'\btacos?\b', r'taquer', r'trompo', r'al pastor', r'tortilla', r'masa harina',
              r'rotisserie', r'pizz', r'neapolit', r'\bdough\b', r'mozzarella',
              r'hamburgueser', r'\bbolas?\b', r'plancha', r'patata', r'fry[- ]?top', r'value meal', r'smashed burger', r'medium[- ]rare',
              r'€', r'\bIVA\b', r'APPCC', r'RD 1021', r'RD 1086', r'Inf26008', r'853/2004', r'\bOrden\b', r'AESAN', r'polares',
              r'(?<![\d.,])175 ?°C', r'(?<![\d.,])70 ?°C(?!\))', r'(?<![\d.,])75 ?°C', r'(?<![\d.,−-])2 ?°C', r'[−-]18 ?°C',
              r'\bkg\b|\bgramos?\b|\bkW\b|\bkWh\b|\bml\b|\bliters?\b|\blitres?\b',
              r'authentic', r'\bcocktails?\b|\bhappy hour\b|\btapas\b',
              r'\bfreelance', r'carnet', r'tarifa plana', r'convenio|\bSMI\b|aut[oó]nomo|Veri\*?factu',
              r'\bunidades\b']
CELDA = {}                                                                         # clave de entrada → (hoja, celda)
for _i in range(NPE):
    for _k in ('pr', 'ut', 'tl', 'gr'):
        CELDA['%s%d' % (_k, _i)] = (HC, '%s%d' % (PC[_i], P_ROW[_k]))
for _k, _f2 in M_C.items():
    CELDA[_k] = (HC, 'B%d' % _f2)
for _k, _f2 in M_A.items():
    CELDA[_k] = (HP, 'B%d' % _f2)
for _k, _f2 in E.items():
    CELDA[_k] = (HE, 'B%d' % _f2)
for _j in range(NPR):
    for _k, _c in (('bo', CBO), ('lq', CLQ), ('pn', CPN), ('ot', COT), ('ra', CRA), ('pvp', CPVP), ('ud', CUD)):
        CELDA['%s%d' % (_k, _j)] = (HE, '%s%d' % (_c, E_P0 + _j))


def celdas_salida(r):
    """(hoja, celda, valor esperado, etiqueta) de todas las salidas."""
    out = []
    for i, c in enumerate(r['Pe']):
        for k in S_C:
            out.append((HC, '%s%d' % (PC[i], S_C[k]), c[k], 'Pe%d.%s' % (i, k)))
    for k in S_D:
        out.append((HC, 'B%d' % S_D[k], r['D'][k], 'D.' + k))
    for k in S_A:
        out.append((HP, 'B%d' % S_A[k], r['A'][k], 'A.' + k))
    out.append((HE, 'B%d' % E_BO, r['bola'], 'bola'))
    out.append((HE, 'B%d' % E_PA, r['pata'], 'pata'))
    for j, q in enumerate(r['Pr']):
        for k, col, _fm in SALP:
            out.append((HE, '%s%d' % (col, E_P0 + j), q[k], '%s[%d]' % (k, j)))
    for k in T:
        out.append((HE, 'B%d' % T[k], r['T'][k], 'T.' + k))
    return out


def _distinto(x, v):
    if isinstance(v, str):
        return x != v
    return not isinstance(x, (int, float)) or abs(float(x) - v) > 0.01


def _checks_logica(nombre, jj, err):
    """Identidades de la lógica sobre calcular(); devuelve cuántas comprobaciones hizo."""
    n = 0
    r = calcular(jj)
    for i, c in enumerate(r['Pe']):
        n += 4
        if abs(c['cku'] * jj['ut%d' % i] * (1 - jj['m']) - jj['pr%d' % i]) > 1e-6:
            err.append('%s: cost per usable lb × usable × (1 − shrink) ≠ price in column %d' % (nombre, i))
        if abs(c['cku_mo'] * jj['ut%d' % i] * (1 - jj['m']) - (jj['pr%d' % i] + jj['tl%d' % i] / 60.0 * jj['hh'])) > 1e-6:
            err.append('%s: cost per usable lb with labor does not add up in column %d' % (nombre, i))
        if abs(c['c_bola'] * jj['OZ'] / jj['gb'] - c['cku_mo']) > 1e-6:
            err.append('%s: cost of a ball ≠ cost per usable lb × ounces ÷ 16 in column %d' % (nombre, i))
        if abs(c['c_dos'] - 2 * c['c_bola']) > 1e-9 or abs(c['c_tres'] - 3 * c['c_bola']) > 1e-9:
            err.append('%s: cost of 2 and 3 balls ≠ 2 and 3 × ball in column %d' % (nombre, i))
    d = r['D']
    n += 3
    if abs(d['bkg'] * jj['gb'] - 16) > 1e-6 or abs(d['bkn'] - d['bkg'] * (1 - jj['m'])) > 1e-6:
        err.append('%s: balls per lb ≠ 16 ÷ oz (or with the shrink)' % nombre)
    if abs(d['carne_dia'] * 16 - jj['cd'] * jj['bc'] * jj['gb']) > 1e-6:
        err.append('%s: lb of beef per day ≠ covers × balls × oz ÷ 16' % nombre)
    a = r['A']
    n += 5
    if abs(a['pat_cong'] * 16 - jj['gc'] * jj['pc']) > 1e-6:
        err.append('%s: frozen fries per portion ≠ oz × $/lb ÷ 16' % nombre)
    if abs(a['kg_bruto'] * jj['uf'] * 16 - jj['gf']) > 1e-6:
        err.append('%s: lb unpeeled × usable share ≠ ounces cut' % nombre)
    if abs(a['rac_dia'] - jj['cd'] * jj['rpc']) > 1e-9:
        err.append('%s: portions per day ≠ covers × portions per cover' % nombre)
    if abs(a['c_cong'] - a['pat_cong'] - a['ace_cong']) > 1e-9:
        err.append('%s: frozen portion ≠ fries + oil' % nombre)
    if abs(a['c_fre_mo'] - a['c_fre'] - a['mo_fre']) > 1e-9:
        err.append('%s: fresh portion with labor ≠ without labor + labor' % nombre)
    for j, q in enumerate(r['Pr']):
        n += 1
        if not q['plat'] < q['llevar'] < q['sala']:
            err.append('%s: marketplace < take-out < dine-in margin fails on product %d' % (nombre, j))
    n += 1
    if abs(r['D']['X'] - (r['Pe'][0]['cku_mo'] - r['Pe'][1]['cku'])) > 1e-9:
        err.append('%s: X ≠ cost per usable lb of the blend − of whole chuck (without labor)' % nombre)
    return n


def autotest(destino):
    from pycel import ExcelCompiler
    err, n = [], 0
    wv = openpyxl.load_workbook(destino, data_only=True)
    wf = openpyxl.load_workbook(destino)
    base = valores_defecto()
    j2 = dict(base, pr0=5.75, ut0=0.98, pr1=4.9, ut1=0.85, tl1=4.0, pr2=6.9, pr3=6.0, ut3=0.9, tl3=2.0, m=0.05, gb=4.0, mc=0.3, hh=20.0, bc=2.0,
              cd=120, cong=OPC[1][0], pc=1.3, gc=6.0, ac=0.06, pf=0.7, gf=7.0, uf=0.7, af=0.12, tf=4.0, hp=20.0, lc=50.0, nc=3, pa=1.6,
              rcm=1200, rpc=0.9, dm=26, tax=0.08, obj=0.32, tol=0.04, sel=2, tp=2, pan=0.6, qs=0.16, cl=0.6, cp=0.9, rec=0.1, com=0.3,
              bo3=3, lq2=1, ra8=2, pvp1=11.5, ud0=100, ud12=300)
    j3 = dict(base, pr0=4.9, ut0=0.97, pr1=4.2, ut1=0.92, tl1=2.0, pr2=6.9, pr3=5.0, m=0.0, gb=2.0, mc=0.2, hh=16.0, bc=1.2, cd=60,
              cong=OPC[2][0], pc=1.0, gc=4.0, ac=0.04, pf=0.45, gf=6.0, uf=0.8, af=0.08, tf=2.0, hp=16.0, lc=35.0, nc=1, pa=1.2, rcm=2000,
              rpc=0.7, dm=22, tax=0.07, obj=0.35, tol=0.1, sel=3, tp=1, pan=0.45, qs=0.12, cl=0.4, cp=0.6, rec=0.2, com=0.35,
              bo0=2, lq1=3, pn4=2, pvp9=7.0, pvp12=17.0, ud5=0, ud12=200)
    # 1. caché y pycel del juego por defecto; pycel de los otros dos con set_value
    xl = ExcelCompiler(filename=destino)
    r0 = calcular(base)
    for hoja, c, v, lab in celdas_salida(r0):
        n += 1
        for et, x in (('pycel', xl.evaluate("'%s'!%s" % (hoja, c))), ('cache', wv[hoja][c].value)):
            if _distinto(x, v):
                err.append('defaults · %s!%s (%s): %s %r ≠ expected %r' % (hoja, c, lab, et, x, v))
    for nombre, jj in (('set 2', j2), ('set 3', j3)):
        r = calcular(jj)
        xl2 = ExcelCompiler(filename=destino)
        for hoja, c, v, lab in celdas_salida(r0):
            xl2.evaluate("'%s'!%s" % (hoja, c))
        for k, v in jj.items():
            if v != base[k]:
                h, c = CELDA[k]
                xl2.set_value("'%s'!%s" % (h, c), v)
        for hoja, c, v, lab in celdas_salida(r):
            n += 1
            x = xl2.evaluate("'%s'!%s" % (hoja, c))
            if _distinto(x, v):
                err.append('%s · %s!%s (%s): pycel %r ≠ expected %r' % (nombre, hoja, c, lab, x, v))
    # 2. identidades de la lógica en los tres juegos
    for nombre, jj in (('defaults', base), ('set 2', j2), ('set 3', j3)):
        n += _checks_logica(nombre, jj, err)
    # identidades sobre la hoja viva (pycel, valores por defecto)
    ev = lambda h, c: xl.evaluate("'%s'!%s" % (h, c))
    for i in range(NPE):
        n += 1
        pr_, ut_ = ev(HC, '%s%d' % (PC[i], P_ROW['pr'])), ev(HC, '%s%d' % (PC[i], P_ROW['ut']))
        m_ = ev(HC, 'B%d' % M_C['m'])
        if abs(ev(HC, '%s%d' % (PC[i], S_C['cku'])) - pr_ / ut_ / (1 - m_)) > 1e-6:
            err.append('Beef & Patty sheet, column %s: cost per usable lb ≠ price ÷ usable ÷ (1 − m)' % PC[i])
    gb_ = ev(HC, 'B%d' % M_C['gb'])
    n += 2
    if abs(ev(HC, 'B%d' % S_D['bkg']) - 16.0 / gb_) > 1e-9:
        err.append('Beef & Patty sheet: balls per lb ≠ 16 ÷ oz')
    if abs(ev(HC, 'B%d' % S_D['bkn']) - 16.0 * (1 - ev(HC, 'B%d' % M_C['m'])) / gb_) > 1e-9:
        err.append('Beef & Patty sheet: balls per lb with the shrink ≠ 16 × (1 − m) ÷ oz')
    # 3. valores con sentido
    sim, dob, cmb = r0['Pr'][J_SIM], r0['Pr'][J_DOB], r0['Pr'][J_COMBO]
    if not 0.18 <= sim['fc'] <= 0.26:
        err.append('single food cost %.1f %% outside 18-26 %%' % (sim['fc'] * 100))
    if not 0.24 <= dob['fc'] <= 0.32:
        err.append('double food cost %.1f %% outside 24-32 %%' % (dob['fc'] * 100))
    if not 0.22 <= r0['T']['fcp'] < 0.34:
        err.append("default weighted food cost %.1f %% is not in 22-34 %% (below the plan's 34 %%)" % (r0['T']['fcp'] * 100))
    for j in (J_COMBO, NPR - 1):
        n += 1
        q = r0['Pr'][j]
        if not q['plat'] < q['llevar'] < q['sala']:
            err.append('product %d: marketplace < take-out < dine-in margin fails at defaults' % j)
    ck = [c['cku'] for c in r0['Pe']]
    if any(abs(a_ - b_) > 0.01 for a_, b_ in zip(ck, (5.41, 5.15, 6.44, 5.93))):
        err.append('default cost per usable lb %r ≠ 5.41 · 5.15 · 6.44 · 5.93' % [round(x, 3) for x in ck])
    D_ = r0['D']
    if abs(D_['bkg'] - 16.0 / 3) > 1e-9 or abs(D_['bkn'] - 16.0 * 0.97 / 3) > 1e-9 or abs(D_['carne_dia'] - 58.5) > 0.01:
        err.append('default balls: 5.33 · 5.17 · 58.5 lb per day (got %r)' % [round(D_[k], 3) for k in ('bkg', 'bkn', 'carne_dia')])
    A_ = r0['A']
    if abs(A_['c_cong'] - 0.459) > 0.005 or abs(A_['c_fre'] - 0.347) > 0.005 or abs(A_['rac_dia'] - 156) > 1e-9:
        err.append('default fries: 0.459 · 0.347 · 156 (got %r)' % [round(A_[k], 3) for k in ('c_cong', 'c_fre', 'rac_dia')])
    if abs(A_['l_mes'] - 322.5) > 0.5:
        err.append('default lb of oil per month 322.5 (got %.2f)' % A_['l_mes'])
    if abs(sim['c_tot'] - 2.005) > 0.01 or abs(dob['c_tot'] - 3.210) > 0.01 or abs(r0['Pr'][2]['c_tot'] - 4.414) > 0.01 \
            or abs(r0['Pr'][3]['c_tot'] - 3.660) > 0.01:
        err.append('default costing: single 2.005 · double 3.210 · triple 4.414 · bacon 3.660 (got %r)'
                   % [round(r0['Pr'][k]['c_tot'], 3) for k in range(4)])
    if abs(cmb['sala'] - 12.07) > 0.02 or abs(cmb['llevar'] - 11.52) > 0.02 or abs(cmb['plat'] - 9.12) > 0.03:
        err.append('default combo margins: 12.07 · 11.52 · 9.12 (got %r)' % [round(cmb[k], 3) for k in ('sala', 'llevar', 'plat')])
    if D_['sem'] != SEM_BLEND or A_['sem'] != SEM_CONG:
        err.append('default verdicts: "%s" and "%s" (expected buy the blend and frozen)' % (D_['sem'], A_['sem']))
    if calcular(dict(base, tl1=0.0))['D']['sem'] != SEM_PICAR:
        err.append('the verdict does not switch to "Grinding in-house pays off" with 0 trimming minutes')
    if calcular(dict(base, tf=0.0))['A']['sem'] != SEM_FRE:
        err.append('the verdict does not switch to "Fresh-cut pays off" with 0 peeling minutes')
    if sorted(q['rank'] for q in r0['Pr'])[0] != 1:
        err.append('the ranks by margin are not coherent')
    # 4. valores expuestos: caché frente a calcular()
    esperados = {'COST_LB_USABLE_BLEND': r0['Pe'][0]['cku'], 'COST_LB_USABLE_CHUCK': r0['Pe'][1]['cku'], 'GRIND_PAYS_OFF_BELOW': D_['X'],
                 'BALLS_PER_LB': D_['bkg'], 'BALLS_PER_LB_NET': D_['bkn'], 'BEEF_PER_DAY_LB': D_['carne_dia'],
                 'COST_SINGLE': sim['c_tot'], 'COST_DOUBLE': dob['c_tot'], 'FC_SINGLE': sim['fc'], 'FC_DOUBLE': dob['fc'],
                 'FC_COMBO': cmb['fc'], 'FC_WEIGHTED': r0['T']['fcp'], 'COST_FROZEN_FRIES_PORTION': A_['c_cong'],
                 'COST_FRESH_FRIES_PORTION': A_['c_fre'], 'OIL_PER_PORTION': A_['ace_cong'], 'OIL_LB_MONTH': A_['l_mes'],
                 'MARGIN_DINEIN_COMBO': cmb['sala'], 'MARGIN_TAKEOUT_COMBO': cmb['llevar'], 'MARGIN_MARKETPLACE_COMBO': cmb['plat']}
    expuestos = valores_expuestos(destino)
    if list(expuestos) != list(esperados):
        err.append('EXPUESTOS ≠ the list of the SPEC: %r' % list(expuestos))
    for nombre, e_ in esperados.items():
        n += 1
        if _distinto(expuestos[nombre], e_):
            err.append('exposed %s: cache %r ≠ %r' % (nombre, expuestos[nombre], e_))
    # 5. formato, protección, DV, fórmulas, textos
    n += 1
    fg = motor.gate_formatos(wf, FICHERO)
    if fg:
        err.append('gate_formatos: %r' % fg[:3])
    dvh = dict((w.title, dict((str(rng), dv) for dv in w.data_validations.dataValidation for rng in dv.sqref.ranges))
               for w in (wf[HC], wf[HP], wf[HE]))
    for hoja, k, tipo in ((HC, 'pr0', 'decimal'), (HC, 'ut1', 'decimal'), (HC, 'm', 'decimal'), (HC, 'cong', 'list'), (HC, 'cd', 'whole'),
                          (HP, 'pc', 'decimal'), (HP, 'nc', 'whole'), (HE, 'obj', 'decimal'), (HE, 'tax', 'decimal'), (HE, 'com', 'decimal'),
                          (HE, 'sel', 'whole'), (HE, 'tp', 'whole'), (HE, 'bo0', 'decimal'), (HE, 'ud0', 'decimal')):
        n += 1
        c = CELDA[k][1]
        dv = dvh[hoja].get(c)
        if dv is None or dv.type != tipo:
            err.append('DV missing or of another type in %s!%s (expected %s)' % (hoja, c, tipo))
    textos = []
    for w in wf.worksheets:
        n += 1
        if w.page_setup.paperSize != 1 or not w.sheet_properties.pageSetUpPr.fitToPage or not w.protection.sheet \
                or w.oddFooter.center.text != FOOTER:
            err.append('%s: Letter + fit + protection + footer' % w.title)
        if w.protection.password:
            err.append('%s: protected with a password' % w.title)
        for row in w.iter_rows():
            for c in row:
                v = c.value
                if v == '':
                    err.append('%s!%s: empty string' % (w.title, c.coordinate))
                if isinstance(v, str) and c.data_type != 'f':
                    textos.append((w.title, c.coordinate, v))
                if isinstance(v, str) and c.data_type == 'f':
                    lits = re.findall(r'"((?:[^"]|"")*)"', v)
                    textos += [(w.title, c.coordinate + ' (formula)', lit) for lit in lits if any(ch.isalpha() for ch in lit)]
                    for lit in lits:
                        if len(lit) > 255:
                            err.append('%s!%s: formula literal of %d characters (> 255)' % (w.title, c.coordinate, len(lit)))
                if c.data_type == 'f':
                    n += 1
                    if wv[w.title][c.coordinate].value is None:
                        err.append('%s!%s: formula without cache' % (w.title, c.coordinate))
                    if len(str(c.value)) >= 200:
                        err.append('%s!%s: formula of %d characters (>= 200)' % (w.title, c.coordinate, len(str(c.value))))
                    sin_refs = re.sub(r'"[^"]*"', '', re.sub(r"'[^']+'!|\$?[A-Z]{1,2}\$?\d+", '', str(c.value)))
                    for num in re.findall(r'(?<![A-Za-z])\d+\.?\d*', sin_refs):
                        if num not in ('0', '1'):
                            err.append('%s!%s: typed number %s in %s' % (w.title, c.coordinate, num, c.value))
                if c.fill.fgColor.rgb in ('00E8F5E9', 'FFE8F5E9') and c.protection.locked:
                    err.append('%s!%s: locked green cell' % (w.title, c.coordinate))
                if c.fill.fgColor.rgb in ('00E8F5E9', 'FFE8F5E9') and c.data_type == 'f':
                    err.append('%s!%s: green cell with a formula' % (w.title, c.coordinate))
                fmt = c.number_format or ''
                if '€' in fmt or '$' in fmt:
                    err.append('%s!%s: currency symbol in the format (D8)' % (w.title, c.coordinate))
        for dv_ in w.data_validations.dataValidation:
            for t in (dv_.error, dv_.errorTitle, dv_.prompt, dv_.promptTitle):
                if t:
                    textos.append((w.title, 'DV %s' % dv_.sqref, t))
    for hoja, ks, mm in ((HC, ('OZ', 'MH'), M_C), (HP, ('OZ', 'MH'), M_A)):
        for k in ks:
            n += 1
            c = wf[hoja]['B%d' % mm[k]]
            rot = wf[hoja]['A%d' % mm[k]].value
            if c.fill.fgColor.rgb not in ('00F2F2F2', 'FFF2F2F2') or c.protection.locked is not True or 'constant' not in rot:
                err.append('gray constant %s badly marked or unlocked' % k)
    p = wf.properties
    for campo in ('title', 'subject', 'keywords', 'description', 'category'):
        textos.append(('docProps', campo, getattr(p, campo) or ''))
    rx_bur = []
    try:
        import datos_bur as D
        rx_bur = [(x[0], x[2] if len(x) > 2 else None) for x in D.RX_BUR_NO]
    except Exception as e:                                                         # noqa: BLE001
        err.append('datos_bur.RX_BUR_NO could not be imported (%s)' % e)
    n_es = n_pz = 0
    for h, c, t in textos:
        n += 1
        r = [x for x in M.restos_espanol(t) if not x[0].startswith('oficio')]
        if r:
            n_es += 1
            err.append('%s!%s: Spanish leftovers %s' % (h, c, r[:1]))
        if M.no_latinos(t):
            err.append('%s!%s: non-Latin characters' % (h, c))
        for rx in PROHIBIDOS:
            if re.search(rx, t, re.I):
                n_pz += 1
                err.append('%s!%s: forbidden /%s/ in %r' % (h, c, rx, t[:60]))
        for rx, exc in rx_bur:
            if re.search(rx, t) and not (exc and re.search(exc, t)):
                n_pz += 1
                err.append('%s!%s: RX_BUR_NO /%s/ in %r' % (h, c, rx, t[:60]))
        if re.search(r'(?i)\bgluten[- ]free', t) and not re.search(r"(?i)\bnever\b|\bnot\b|\bno\b|can't|cannot|don't", t):
            err.append('%s!%s: gluten-free without a negation' % (h, c))
        if re.search(r'(?i)\bvalue\b.*\$\d|\bsave \$|total value|\bwas \$|\d+\s?% off\b', t):
            err.append('%s!%s: price anchor' % (h, c))
    todo = '\n'.join(t for _h, _c, t in textos)
    for frase in ('16 oz per lb', '60 minutes per hour', TAX_ROT, '155 °F (68 °C)', '17 seconds', '3-401.11(A)(2)', '3-603.11', '3-501.16', '3-302.11',
                  '9 CFR 320.1(b)(4)', 'Prop 65', 'TPM', SEM_PICAR, SEM_BLEND, SEM_FRE, SEM_CONG, FC_OK, FC_VIG, FC_ALTO,
                  'Kit estimate, editable', "Kit estimate (Oct 2026): replace with your distributor's price", 'no source: ',
                  'Burger Restaurant Financial Projections', 'Burger Restaurant Startup Checklist', 'Burger Restaurant Business Plan',
                  'Single smash burger (1 ball)', 'Double smash burger (2 balls)', 'Combo: double + fries + soft drink',
                  'good practice', 'Mustard is not a major allergen', 'Never promise "gluten-free"', 'the crust is not a thermometer',
                  'the crust softens in about 15 minutes'):
        n += 1
        if frase not in todo:
            err.append('missing text %r' % frase)
    n += 1
    if M.FIRMA_EN not in todo or VERSION_LINEA not in todo or M.AVISO_EN not in todo:
        err.append('missing signature, version line or notice')
    if wf[HE]['A5'].value != TAX_ROT or wf[HE]['B5'].value != 0:
        err.append('Costing!B5 is not the sales tax = 0 %')
    if wf.sheetnames != [HC, HP, HE, HI]:
        err.append('sheet names %r' % wf.sheetnames)
    if NPR != 14:
        err.append('not 14 products')
    if p.creator != 'AI Chef Pro' or p.description != URL or 'AI Chef Pro' not in (p.keywords or '') \
            or not (p.title or '').startswith(TITULO):
        err.append('docProps: creator=%r title=%r description=%r keywords=%r' % (p.creator, p.title, p.description, p.keywords))
    for x in err[:80]:
        print('ERROR', x)
    print('autotest smash burger patty, fries & costing calculator: %d checks (pycel + cache at defaults; pycel in 2 more sets; identities) '
          'over %d texts · %s' % (n, len(textos), 'GREEN' if not err else 'RED (%d errors)' % len(err)))
    print('  beef ($/usable lb): ' + ' · '.join('%s %.3f' % (CARNES[i][0], c['cku']) for i, c in enumerate(r0['Pe'])))
    print('  grinding: labor $%.3f per usable lb · pays off up to $%.3f per usable lb · %s' % (D_['lab_kg'], D_['X'], D_['sem']))
    print('  balls: %.2f per lb raw · %.2f with the shrink · %.1f balls/day · %.2f lb of beef/day' % (D_['bkg'], D_['bkn'], D_['bdia'], D_['carne_dia']))
    print('  fries: frozen $%.4f (oil %.4f) · fresh $%.4f without labor, $%.4f with it · fresh-cut pays off up to $%.3f/lb (labor %.3f) · %s'
          % (A_['c_cong'], A_['ace_cong'], A_['c_fre'], A_['c_fre_mo'], A_['X'], A_['mo_kg'], A_['sem']))
    print('  fries: %.0f portions/day · %.1f lb frozen/day · %.1f lb fresh/day · %.1f lb of oil/month' % (
        A_['rac_dia'], A_['kg_cong_dia'], A_['kg_fre_dia'], A_['l_mes']))
    for nombre, (h, c) in EXPUESTOS.items():
        print('  %-28s %s!%-5s %s' % (nombre, h, c, expuestos[nombre]))
    for j, p_ in enumerate(PRODS):
        q = r0['Pr'][j]
        print('  %-42s cost $%.3f · food cost %5.1f %% · %-14s · margin dine-in %.2f / take-out %.2f / marketplace %.2f · rank %d · %d/week'
              % (p_[0], q['c_tot'], q['fc'] * 100, q['sem'], q['sala'], q['llevar'], q['plat'], q['rank'], p_[7]))
    print('  weighted: sales $%.2f · cost $%.2f · food cost %.2f %% · %s' % (r0['T']['ing'], r0['T']['cos'], r0['T']['fcp'] * 100, r0['T']['semp']))
    print('  combo: marketplace margin vs dine-in %.1f %%' % ((cmb['plat'] / cmb['sala'] - 1) * 100))
    return 1 if err else 0


if __name__ == '__main__':
    ap = argparse.ArgumentParser()
    ap.add_argument('--salida', required=True, help='output folder')
    ap.add_argument('--autotest', action='store_true')
    a = ap.parse_args()
    destino = os.path.join(a.salida, FICHERO)
    if not a.autotest and generar(destino):
        sys.exit('ERROR: fallos de pycel en inject_cache')
    sys.exit(autotest(destino))
