#!/usr/bin/env python3
"""
calculadora_pol.py — Rotisserie Chicken Yield & Costing Calculator, calculadora del Rotisserie Chicken Restaurant Business Plan Kit
(EN, $39). Es la edición EE. UU. de `../plan-polleria/calculadora_pol_es.py`: MISMAS hojas (`Chicken Yield` · `Rotisserie` ·
`Costing` · `Instructions`), filas, fórmulas, helpers, `inject_cache` y autotest (pycel + caché = `calcular()`, 3 juegos), con el
patrón de `../business-plans-en-12/calculadora_taq.py`. Cambios EN:

  (1) pollo en libras ($/lb, WOG) y porciones en ONZAS, con la constante gris «16 oz per lb» (el gate prohíbe números tecleados
      en fórmulas salvo 0/1); °F en el texto de las instrucciones (165 °F cocción, 135 °F hot holding), nunca como entrada;
  (2) coste del pollo entero = (3.0 lb × $1.75/lb + marinado $0.45) × (1 + 5 % de merma de display) = $5.985 ≈ $5.99;
  (3) hoja `Rotisserie`: dos rotisseries de gas (entrada «Rotisseries» = 2) × 8 spits × 6 pollos = 96 por carga; el gas se
      calcula con BTU/h por rotisserie y $/therm (constante gris «BTU per therm» = 100,000);
  (4) hoja `Costing`: whole · half + 2 sides · quarter + 2 sides · family meal · sandwich/bowl · salchipapas · large yuca fries ·
      maduros · soda; sales tax en carta = 0 %, comisión del marketplace 25 %; margen dine-in · pickup & pre-orders · marketplace;
  (5) US Letter, formatos sin símbolo de moneda (los rótulos llevan «($)»), docProps en inglés, versión del libro EN.

Valores que otros scripts leen de la CACHÉ del xlsx (dict `EXPUESTOS`, función `valores_expuestos(ruta)`):
    BIRDS_PER_LOAD     Rotisserie!B<carga>      96 pollos (2 × 8 spits × 6)
    SUNDAY_CAPACITY    Rotisserie!B<cap_dom>    192 pollos antes de la hora límite (2 cargas × 96)
    COOKED_YIELD       Chicken Yield!B<rend>    0.72
    COST_PER_BIRD      Chicken Yield!B<coste>   $5.985 ≈ $5.99
    FAMILY_MEAL_FC     Costing!G<fc>            ≈ 29.6 % (family meal, sobre el precio antes de tax)
Corre en el VPS (openpyxl + pycel):

    /root/venv-guias/bin/python calculadora_pol.py --salida <dir>             # genera + inject_cache + autotest
    /root/venv-guias/bin/python calculadora_pol.py --salida <dir> --autotest

Sesión Claude Code, 29-oct-2026.
"""
import argparse
import math
import os
import re
import sys
from collections import OrderedDict

import openpyxl
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill, Protection
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

AQUI = os.path.dirname(os.path.abspath(__file__))
PD = os.path.dirname(AQUI)
os.environ.setdefault('BP_CONJUNTO', 'pol')
sys.path.insert(0, os.path.join(PD, 'planes-v2_0'))
sys.path.insert(0, os.path.join(PD, 'business-plans-en'))
sys.path.insert(0, PD)
sys.path.insert(0, AQUI)
import motor                                                                       # noqa: E402  (solo gate_formatos)
import mapas as M                                                                  # noqa: E402  (detectores y textos fijos)

_PL = M.PLANES['pol']
SLUG = _PL['slug']
PRODUCTO = _PL['producto']
URL = _PL['url']
FICHERO = 'rotisserie-chicken-yield-costing-calculator.xlsx'
TITULO = 'Rotisserie Chicken Yield & Costing Calculator (Chicken Cost, Rotisserie Loads and Plate Costing)'
KEYWORDS = _PL['keywords']
VERSION, MES = M.VERSION, M.MES_VERSION
VERSION_LINEA = 'Version %s · %s · %s · info@aichef.pro' % (VERSION, MES, URL)
FOOTER = M.FOOTER_EN
HP, HA, HE, HI = 'Chicken Yield', 'Rotisserie', 'Costing', 'Instructions'
VERDE = PatternFill('solid', fgColor='E8F5E9')
GRIS = PatternFill('solid', fgColor='F2F2F2')
ROJO = PatternFill('solid', bgColor='F8CBAD', fgColor='F8CBAD')
AMBAR = PatternFill('solid', bgColor='FFE699', fgColor='FFE699')
OK = PatternFill('solid', bgColor='C6E0B4', fgColor='C6E0B4')
FUENTE = 'Calibri'
USD, ENT, PCT, DEC, OZF, USD3 = '#,##0.00', '#,##0', '0.0%', '#,##0.0', '#,##0.0', '#,##0.000'
FC_OK, FC_VIG, FC_ALTO = 'On target', 'Watch the cost', 'Above target'
SEM_OK, SEM_KO = 'Covered', 'Short at the peak'
EST = ' Kit estimate, editable.'
KIT_PRECIO = "Kit estimate (Oct 2026): replace with your distributor's price"

# ---------------------------------------------------------------------------------------------------------------------
# Datos (F1-research-us-delta §4-§5). Importes antes de sales tax. Todo «Kit estimate (Oct 2026)», editable.
# ---------------------------------------------------------------------------------------------------------------------
#: hoja Chicken Yield: (clave, rótulo, formato, valor por defecto, constante gris, nota)
ENT_P = [
    ('P', 'Average whole chicken weight (lb, WOG)', '#,##0.00', 3.0, False,
     'Broiler/fryer of 2.5 to 4.5 lb WOG (without giblets). Weigh your birds at receiving.' + EST),
    ('Pk', 'Whole chicken price ($/lb, before tax)', USD, 1.75, False,
     'Delivered price of fresh whole chickens. The USDA wholesale composite ran about $1.17/lb in August 2026 and a distributor adds freight and margin. '
     "Type your supplier's price." + EST),
    ('S', 'Marinade per bird ($, before tax)', USD, 0.45, False,
     'Soy sauce, beer, garlic, cumin, aji panca, huacatay and vinegar, marinated 24 to 48 hours.' + EST),
    ('hw', 'Display and holding waste (% of the bird cost)', PCT, 0.05, False,
     'Birds that pass their hot-holding window. Kit target: 5 % or less of the birds you load.' + EST),
    ('r', 'Cooked yield (cooked weight ÷ raw weight)', PCT, 0.72, False,
     'What is left of the bird after roasting, bone-in. No manufacturer figure is published. Measure it with 5 birds from the same load (see the "Instructions" sheet).' + EST),
    ('ref', 'Benchmark: kit range for a rotisserie bird (midpoint of 70 to 75 %)', PCT, 0.725, True,
     'Fixed benchmark: the kit works with a cooked yield of 70 to 75 %. Not editable.'),
    ('pp', 'Guests who share one bird', ENT, 3, False, 'Whole number greater than 0. For a main plate with two sides it is usually 3.' + EST),
    ('Np', 'Guests you want to serve (pre-order or party)', ENT, 40, False, 'Whole number greater than 0. It tells you how many birds to order.'),
    ('Ns', 'Birds left unsold at close', DEC, 2.0, False, 'Whole birds that stay unsold when you close. Count them; default 2.' + EST),
    ('a', 'Share of the leftover you reuse (%)', PCT, 0.50, False,
     'Next-day sandwich or bowl meat, after correct cooling (135 to 70 °F within 2 hours and to 41 °F within 6 hours in total). The part you throw away costs money.' + EST),
    ('OZ', 'Ounces per lb (constant: 16 oz per lb)', ENT, 16, True, 'Constant: 1 lb = 16 oz. Not editable. It converts the cooked pounds to ounces.'),
    ('M2', 'Halves per bird (constant)', ENT, 2, True, 'Constant: a bird is cut into 2 halves. Not editable.'),
    ('M4', 'Quarters per bird (constant)', ENT, 4, True, 'Constant: a bird is cut into 4 quarters. Not editable.'),
]
SAL_P = [
    ('asado', 'COOKED WEIGHT OF ONE BIRD (lb)', '#,##0.000', '{P}*{r}', 'Raw weight × cooked yield.'),
    ('merma_kg', 'Cooking loss (lb)', '#,##0.000', '{P}-{asado}', 'Water, fat and skin that leave the rotisserie.'),
    ('merma_pct', 'Cooking loss (%)', PCT, '{merma_kg}/{P}', 'Loss ÷ raw weight.'),
    ('rend', 'COOKED YIELD (%)', PCT, '{asado}/{P}', 'Cooked weight ÷ raw weight.'),
    ('vs_ref', 'Your cooked yield against the kit benchmark (%)', '0%', '{rend}/{ref}', 'Below 100 %: you lose more than the kit midpoint; weigh again.'),
    ('coste', 'COST PER WHOLE BIRD ($)', USD3, '({P}*{Pk}+{S})*(1+{hw})', 'Weight × price per lb + marinade, plus the holding waste. It feeds the costing.'),
    ('coste_kg', 'COST PER COOKED LB ($)', USD, '{coste}/{asado}', 'Bird cost ÷ cooked pounds.'),
    ('g_entero', 'Cooked ounces of a whole bird (oz)', OZF, '{asado}*{OZ}', 'Cooked weight × 16.'),
    ('g_medio', 'Cooked ounces of a half chicken (oz)', OZF, '{g_entero}/{M2}', 'Whole cooked bird ÷ 2.'),
    ('g_cuarto', 'Cooked ounces of a quarter chicken (oz)', OZF, '{g_entero}/{M4}', 'Whole cooked bird ÷ 4.'),
    ('c_medio', 'Cost of a half chicken ($)', USD3, '{coste}/{M2}', 'Bird cost ÷ 2.'),
    ('c_cuarto', 'Cost of a quarter chicken ($)', USD3, '{coste}/{M4}', 'Bird cost ÷ 4.'),
    ('pollos_N', 'BIRDS FOR THE GUESTS (birds)', ENT, 'ROUNDUP({Np}/{pp},0)', 'Guests ÷ guests per bird, rounded up.'),
    ('kg_canal', 'Raw pounds you have to buy (lb)', '#,##0.00', '{pollos_N}*{P}', 'Birds × raw weight.'),
    ('c_merma', 'Cost of the leftover you throw away ($)', USD, '{Ns}*(1-{a})*{coste}', 'Leftover birds × share you do not reuse × bird cost.'),
]

#: hoja Rotisserie
ENT_S = [
    ('R', 'Rotisseries', ENT, 2, False, 'Whole number greater than 0. The kit case runs two gas rotisseries of 8 spits each.'),
    ('E', 'Spits per rotisserie', ENT, 8, False, 'Whole number greater than 0. Type your model (8-spit units are common).'),
    ('p', 'Birds per spit', ENT, 6, False, 'Spec sheets give about 6 chickens per spit on an 8-spit gas unit (48 birds). Whole number greater than 0.'),
    ('t', 'Cook minutes of one load', DEC, 80.0, False,
     'NO PUBLISHED FIGURE: no spec sheet gives the time of a 3 lb bird. The kit works with 75 to 90 minutes. Time your own unit (see the "Instructions" sheet).' + EST),
    ('c', 'Load, unload and clean minutes between loads', DEC, 15.0, False, 'Skewering the next load while the last one comes out.' + EST),
    ('hh0', 'First load: hour (24-hour clock)', ENT, 9, False, 'Whole hour from 0 to 23. With 9 and 0 minutes, the first load goes on at 9:00 AM.' + EST),
    ('mm0', 'First load: minutes', ENT, 0, False, 'Minutes from 0 to 59.'),
    ('hh1', 'Deadline to have the birds cooked: hour (24-hour clock)', ENT, 13, False,
     'The Sunday lunch peak runs until about 1:30 PM. Whole hour from 0 to 23.' + EST),
    ('mm1', 'Deadline to have the birds cooked: minutes', ENT, 30, False, 'Minutes from 0 to 59.'),
    ('D', 'Birds you sell on a normal Sunday', ENT, 68, False, 'About 1.6 to 1.8 times a weekday in the kit case (65 to 70 birds). Replace it with yours.' + EST),
    ('hp', 'Share of the Sunday sold in the peak window (%)', PCT, 0.70, False, 'Between about 11:30 AM and 1:30 PM.' + EST),
    ('ag', 'Increase on a holiday or peak day (%)', PCT, 1.00, False,
     'Over a normal Sunday: Mother\'s Day or a Super Bowl Sunday reach 120 to 150 birds in the kit case.' + EST),
    ('btu', 'Gas input per rotisserie (BTU/h)', ENT, 177500, False,
     'A gas 8-spit unit is rated about 177,500 BTU/h at full burn; the real average is lower. Type your unit\'s rating.' + EST),
    ('th', 'Gas price ($ per therm, before tax)', USD, 1.30, False, 'Check your gas bill: the price per therm is on it.' + EST),
    ('MH', 'Minutes per hour (constant)', ENT, 60, True, 'Constant: 1 hour = 60 minutes. Not editable.'),
    ('BT', 'BTU per therm (constant)', ENT, 100000, True, 'Constant: 1 therm = 100,000 BTU. Not editable.'),
]
SAL_S = [
    ('carga', 'BIRDS PER LOAD', ENT, '{R}*{E}*{p}', 'Rotisseries × spits × birds per spit.'),
    ('carga_una', 'Birds per load, each rotisserie', ENT, '{E}*{p}', 'Spits × birds per spit.'),
    ('ciclo', 'Minutes of a full cycle', DEC, '{t}+{c}', 'Cook + load and unload.'),
    ('tandas_h', 'Loads per hour', '#,##0.00', '{MH}/{ciclo}', '60 minutes ÷ minutes of the cycle.'),
    ('cap_h', 'Capacity per hour (birds)', DEC, '{carga}*{tandas_h}', 'Birds per load × loads per hour.'),
    ('min_disp', 'Minutes between the first load and the deadline', ENT, '({hh1}-{hh0})*{MH}+({mm1}-{mm0})', 'If it is negative, the deadline is before the first load.'),
    ('tandas', 'LOADS BEFORE THE DEADLINE', ENT, 'INT(MAX(0,{min_disp})/{ciclo})', 'Full cycles that fit in that time.'),
    ('cap_dom', 'SUNDAY CAPACITY BEFORE THE DEADLINE (birds)', ENT, '{carga}*{tandas}', 'Birds per load × loads.'),
    ('dem_punta', 'Birds sold in the peak window of a normal Sunday', DEC, '{D}*{hp}', 'Sunday birds × share sold in the peak window.'),
    ('cargas_dom', 'Loads a normal Sunday needs (all day)', ENT, 'ROUNDUP({D}/{carga},0)', 'Sunday birds ÷ birds per load, rounded up.'),
    ('cargas_punta', 'Loads the peak window of a normal Sunday needs', ENT, 'ROUNDUP({dem_punta}/{carga},0)', 'Peak birds ÷ birds per load, rounded up.'),
    ('balance', 'Surplus (+) or shortfall (−) at the peak (birds)', DEC, '{cap_dom}-{dem_punta}', 'Capacity before the deadline minus the peak birds.'),
    ('sem', 'Peak traffic light (normal Sunday)', 'General', 'IF({balance}>=0,"%s","%s")' % (SEM_OK, SEM_KO), 'Covered if the capacity reaches the peak birds.'),
    ('d_ago', 'Birds on a holiday or peak day', ENT, 'ROUNDUP({D}*(1+{ag}),0)', 'Normal Sunday plus the increase, rounded up.'),
    ('cargas_ago', 'Loads a holiday or peak day needs', ENT, 'ROUNDUP({d_ago}/{carga},0)', 'Peak-day birds ÷ birds per load, rounded up.'),
    ('balance_ago', 'Surplus (+) or shortfall (−) at the peak of a peak day (birds)', DEC, '{cap_dom}-{d_ago}*{hp}', 'Capacity minus the peak-day peak.'),
    ('sem_ago', 'Peak traffic light (holiday or peak day)', 'General', 'IF({balance_ago}>=0,"%s","%s")' % (SEM_OK, SEM_KO), 'Covered if the capacity reaches the peak of a peak day.'),
    ('tarde_min', 'Latest first load (minutes since 00:00)', ENT, '{hh1}*{MH}+{mm1}-{cargas_punta}*{ciclo}', 'Deadline minus the peak loads times the cycle.'),
    ('tarde_h', 'Latest first load: hour', ENT, 'INT({tarde_min}/{MH})', 'If you load later, you miss the peak.'),
    ('tarde_m', 'Latest first load: minutes', ENT, 'MOD({tarde_min},{MH})', ''),
    ('horas_gas', 'Hours of rotisserie burning on a normal Sunday', '#,##0.00', '{cargas_dom}*{ciclo}/{MH}', 'Loads × cycle ÷ 60.'),
    ('gas_dia', 'Gas of a normal Sunday ($)', USD, '{btu}*{R}*{horas_gas}/{BT}*{th}', 'BTU/h × rotisseries × hours ÷ BTU per therm × price per therm.'),
    ('gas_pollo', 'GAS PER COOKED BIRD ($)', USD3, '{gas_dia}/({cargas_dom}*{carga})', 'Gas of the day ÷ birds that fit in the day\'s loads.'),
]

#: hoja Costing
ESC_IN = [
    ('tax', 'Sales tax included in your menu prices (%)', PCT, 0.0,
     'Typically 0 %: a US menu is quoted before tax and the tax is added at the register. Rates vary by state and city; ask your accountant.'),
    ('obj', 'Target food cost (% of the price before tax)', PCT, 0.30, 'Rotisserie chicken reference: food and beverage at 30 % or less.' + EST),
    ('tol', 'Watch margin over the target (points)', PCT, 0.05, 'Up to this margin above the target, the traffic light says "Watch the cost". The whole bird tends to land here.'),
    ('merm', 'Production waste on the sides (%)', PCT, 0.0,
     'Default 0: the bird already carries its holding waste on the "Chicken Yield" sheet. Type a percentage if you throw away fries, salad or rice.' + EST),
    ('env', 'Packaging per takeout order ($, before tax)', USD, 0.60, 'Thermal bag, tray and sauce cups of one order.' + EST),
    ('com', 'Delivery marketplace commission (%)', PCT, 0.25, 'Over the price before tax. Check your contract with each marketplace.' + EST),
    ('rec', 'Menu markup on the marketplace (%)', PCT, 0.15, 'How much you raise the price on the marketplace to offset the commission.' + EST),
]
#: onzas: (ingrediente, $/lb, origen del precio)
YU, FR, RB, SA, MA, AV, SS = ('Yuca fries (frozen, with fry oil)', 'French fries (frozen, with fry oil)', 'Rice & beans (cooked)',
                              'House salad', 'Maduros (fried sweet plantains)', 'Aji verde (Peruvian green sauce)', 'Sliced sausage (for salchipapas)')
ING_G = [
    (YU, 2.97, 'Kit estimate (Oct 2026): frozen yuca about $2.75/lb plus fry oil. Replace with your distributor\'s price'),
    (FR, 2.20, KIT_PRECIO),
    (RB, 1.20, KIT_PRECIO),
    (SA, 2.40, KIT_PRECIO),
    (MA, 3.00, KIT_PRECIO),
    (AV, 2.00, KIT_PRECIO),
    (SS, 4.50, KIT_PRECIO),
]
#: unidades: (ingrediente, $/ud o None si es el coste del pollo de la hoja Chicken Yield, origen)
UPOL, UPAN, UBEB = 'Rotisserie chicken (fraction of a whole bird)', 'Bread roll (ea)', 'Peruvian or other soda, can or bottle (ea)'
ING_U = [
    (UPOL, None, 'linked from the Chicken Yield sheet (cost per whole bird)'),
    (UPAN, 0.40, KIT_PRECIO),
    (UBEB, 0.70, KIT_PRECIO),
]
#: columna → (precio de carta, {ingrediente en oz: oz}, {unidades: cantidad})
RECETAS = OrderedDict([
    ('Whole chicken', (19.00, {}, {UPOL: 1})),
    ('Half chicken + 2 sides', (17.50, {YU: 7.5, SA: 4, AV: 2}, {UPOL: 0.5})),
    ('Quarter chicken + 2 sides', (13.00, {YU: 7.5, SA: 4, AV: 2}, {UPOL: 0.25})),
    ('Family meal (whole chicken + 2 large sides)', (34.00, {YU: 14.4, RB: 12, AV: 4}, {UPOL: 1})),
    ('Chicken sandwich or bowl', (13.50, {RB: 6, SA: 3, AV: 1}, {UPOL: 0.25, UPAN: 1})),
    ('Salchipapas', (10.50, {FR: 10, SS: 4, AV: 2}, {})),
    ('Large yuca fries (side)', (8.00, {YU: 12, AV: 2}, {})),
    ('Maduros (side)', (4.50, {MA: 5}, {})),
    ('Peruvian or other soda', (3.25, {}, {UBEB: 1})),
])
NP, NG, NU = len(RECETAS), len(ING_G), len(ING_U)
PCOL = [get_column_letter(4 + j) for j in range(NP)]                              # D..L
ORI = get_column_letter(4 + NP)                                                    # M: origen del precio
J_ENTERO, J_FAMILIA = 0, 3

# ---------------------------------------------------------------------------------------------------------------------
# Disposición de las hojas (filas)
# ---------------------------------------------------------------------------------------------------------------------
def _filas(ent, sal, h0=5):
    m = OrderedDict((k[0], h0 + 1 + i) for i, k in enumerate(ent))
    ra = h0 + 1 + len(ent) + 1
    s = OrderedDict((k[0], ra + 1 + i) for i, k in enumerate(sal))
    return m, ra, s


M_P, M_RP, S_P = _filas(ENT_P, SAL_P)                                              # hoja Chicken Yield
M_S, M_RS, S_S = _filas(ENT_S, SAL_S)                                              # hoja Rotisserie
E = OrderedDict((k[0], 5 + i) for i, k in enumerate(ESC_IN))
E_CAB = 5 + len(ESC_IN) + 1
E_PVP = E_CAB + 1
E_G0 = E_PVP + 1
E_G1 = E_G0 + NG - 1
E_UH = E_G1 + 1
E_U0 = E_UH + 1
E_U1 = E_U0 + NU - 1
E_R = E_U1 + 2
RES = [
    ('peso', 'Weight served on the plate (oz)', OZF),
    ('c_pollo', 'Cost of the chicken ($)', USD3),
    ('c_acomp', 'Cost of sides, sauces, bread and drinks, before waste ($)', USD3),
    ('c_tot', 'COST WITH WASTE ON THE SIDES ($)', USD3),
    ('pvp_sin', 'Menu price before tax ($)', USD),
    ('fc', 'FOOD COST (%)', PCT),
    ('sem', 'Traffic light against the target', 'General'),
    ('margen', 'Gross margin ($, price before tax − cost)', USD),
    ('pvp_sug', 'Menu price with tax needed to hit the target ($)', USD),
]
R = OrderedDict((k, E_R + 1 + i) for i, (k, *_r) in enumerate(RES))
E_CH = E_R + len(RES) + 2
CANAL = [
    ('sala_e', 'Dine-in: margin ($)', USD), ('sala_p', 'Dine-in: margin (%)', PCT),
    ('mostr_e', 'Pickup & pre-orders: margin ($, after packaging)', USD), ('mostr_p', 'Pickup & pre-orders: margin (%)', PCT),
    ('plat_pv', 'Marketplace: price before tax with markup ($)', USD),
    ('plat_e', 'Marketplace: margin ($, after packaging and commission)', USD), ('plat_p', 'Marketplace: margin (%)', PCT),
    ('rank', 'Rank by pickup margin (1 = the one that earns the most)', ENT),
]
CH = OrderedDict((k, E_CH + 1 + i) for i, (k, *_r) in enumerate(CANAL))
E_FIN = E_CH + len(CANAL) + 2

EXPUESTOS = OrderedDict([
    ('BIRDS_PER_LOAD', (HA, 'B%d' % S_S['carga'])),
    ('SUNDAY_CAPACITY', (HA, 'B%d' % S_S['cap_dom'])),
    ('COOKED_YIELD', (HP, 'B%d' % S_P['rend'])),
    ('COST_PER_BIRD', (HP, 'B%d' % S_P['coste'])),
    ('FAMILY_MEAL_FC', (HE, '%s%d' % (PCOL[J_FAMILIA], R['fc']))),
])


def valores_defecto():
    """Todas las entradas por defecto en un solo dict plano (claves = las de las celdas verdes)."""
    v = OrderedDict()
    for ent in (ENT_P, ENT_S):
        for k, _r, _f2, val, _c, _n in ent:
            v[k] = val
    for k, _r, _f2, val, _n in ESC_IN:
        v[k] = val
    for i, (_n, p, _o) in enumerate(ING_G):
        v['p%d' % i] = p
    for i, (_n, p, _o) in enumerate(ING_U):
        if p is not None:
            v['up%d' % i] = p
    for j, (_nom, (pvp, gr, ud)) in enumerate(RECETAS.items()):
        v['pvp%d' % j] = pvp
        for i, (ni, *_r) in enumerate(ING_G):
            v['g%d_%d' % (i, j)] = float(gr.get(ni, 0))
        for i, (ni, *_r) in enumerate(ING_U):
            v['u%d_%d' % (i, j)] = float(ud.get(ni, 0))
    return v


# ---------------------------------------------------------------------------------------------------------------------
# Cálculo en Python (espejo de las fórmulas) — fuente del autotest
# ---------------------------------------------------------------------------------------------------------------------
def calcular(v):
    r = OrderedDict()
    p = OrderedDict()
    p['asado'] = v['P'] * v['r']
    p['merma_kg'] = v['P'] - p['asado']
    p['merma_pct'] = p['merma_kg'] / v['P']
    p['rend'] = p['asado'] / v['P']
    p['vs_ref'] = p['rend'] / v['ref']
    p['coste'] = (v['P'] * v['Pk'] + v['S']) * (1 + v['hw'])
    p['coste_kg'] = p['coste'] / p['asado']
    p['g_entero'] = p['asado'] * v['OZ']
    p['g_medio'] = p['g_entero'] / v['M2']
    p['g_cuarto'] = p['g_entero'] / v['M4']
    p['c_medio'] = p['coste'] / v['M2']
    p['c_cuarto'] = p['coste'] / v['M4']
    p['pollos_N'] = math.ceil(v['Np'] / v['pp'])
    p['kg_canal'] = p['pollos_N'] * v['P']
    p['c_merma'] = v['Ns'] * (1 - v['a']) * p['coste']
    r['P'] = p
    s = OrderedDict()
    s['carga'] = v['R'] * v['E'] * v['p']
    s['carga_una'] = v['E'] * v['p']
    s['ciclo'] = v['t'] + v['c']
    s['tandas_h'] = v['MH'] / s['ciclo']
    s['cap_h'] = s['carga'] * s['tandas_h']
    s['min_disp'] = (v['hh1'] - v['hh0']) * v['MH'] + (v['mm1'] - v['mm0'])
    s['tandas'] = math.floor(max(0, s['min_disp']) / s['ciclo'])
    s['cap_dom'] = s['carga'] * s['tandas']
    s['dem_punta'] = v['D'] * v['hp']
    s['cargas_dom'] = math.ceil(v['D'] / s['carga'])
    s['cargas_punta'] = math.ceil(s['dem_punta'] / s['carga'])
    s['balance'] = s['cap_dom'] - s['dem_punta']
    s['sem'] = SEM_OK if s['balance'] >= 0 else SEM_KO
    s['d_ago'] = math.ceil(v['D'] * (1 + v['ag']))
    s['cargas_ago'] = math.ceil(s['d_ago'] / s['carga'])
    s['balance_ago'] = s['cap_dom'] - s['d_ago'] * v['hp']
    s['sem_ago'] = SEM_OK if s['balance_ago'] >= 0 else SEM_KO
    s['tarde_min'] = v['hh1'] * v['MH'] + v['mm1'] - s['cargas_punta'] * s['ciclo']
    s['tarde_h'] = math.floor(s['tarde_min'] / v['MH'])
    s['tarde_m'] = s['tarde_min'] - v['MH'] * math.floor(s['tarde_min'] / v['MH'])
    s['horas_gas'] = s['cargas_dom'] * s['ciclo'] / v['MH']
    s['gas_dia'] = v['btu'] * v['R'] * s['horas_gas'] / v['BT'] * v['th']
    s['gas_pollo'] = s['gas_dia'] / (s['cargas_dom'] * s['carga'])
    r['S'] = s
    r['pg'] = [v['p%d' % i] for i in range(NG)]
    r['pu'] = [p['coste']] + [v['up%d' % i] for i in range(1, NU)]
    cols = []
    for j in range(NP):
        gs = [v['g%d_%d' % (i, j)] for i in range(NG)]
        us = [v['u%d_%d' % (i, j)] for i in range(NU)]
        q = OrderedDict()
        q['peso'] = sum(gs) + us[0] * p['g_entero']
        q['c_pollo'] = us[0] * r['pu'][0]
        q['c_acomp'] = sum(g_ * pr for g_, pr in zip(gs, r['pg'])) / v['OZ'] + sum(u_ * pr for u_, pr in zip(us[1:], r['pu'][1:]))
        q['c_tot'] = q['c_pollo'] + q['c_acomp'] * (1 + v['merm'])
        q['pvp_sin'] = v['pvp%d' % j] / (1 + v['tax'])
        q['fc'] = q['c_tot'] / q['pvp_sin']
        q['sem'] = FC_OK if q['fc'] <= v['obj'] else (FC_VIG if q['fc'] <= v['obj'] + v['tol'] else FC_ALTO)
        q['margen'] = q['pvp_sin'] - q['c_tot']
        q['pvp_sug'] = q['c_tot'] / v['obj'] * (1 + v['tax'])
        q['sala_e'] = q['margen']
        q['sala_p'] = q['sala_e'] / q['pvp_sin']
        q['mostr_e'] = q['sala_e'] - v['env']
        q['mostr_p'] = q['mostr_e'] / q['pvp_sin']
        q['plat_pv'] = q['pvp_sin'] * (1 + v['rec'])
        q['plat_e'] = q['plat_pv'] * (1 - v['com']) - q['c_tot'] - v['env']
        q['plat_p'] = q['plat_e'] / q['plat_pv']
        cols.append(q)
    for q in cols:
        q['rank'] = 1 + sum(1 for o in cols if o['mostr_e'] > q['mostr_e'])
    for j, q in enumerate(cols):
        r['p%d' % j] = q
    return r


# ---------------------------------------------------------------------------------------------------------------------
# Helpers (del Taco EN)
# ---------------------------------------------------------------------------------------------------------------------
def _f(ws, coord, v, b=False, sz=11, i=False, fmt=None, wrap=False, fill=None, locked=True):
    c = ws[coord]
    c.value = v
    c.font = Font(name=FUENTE, size=sz, bold=b, italic=i)
    if fmt:
        c.number_format = fmt
    c.alignment = Alignment(wrap_text=wrap, vertical='top')
    if fill:
        c.fill = fill
    c.protection = Protection(locked=locked)
    return c


def _nota(ws, fila, texto, desde='E', hasta='F', chars=80.0):
    if not texto:
        return
    _f(ws, '%s%d' % (desde, fila), texto, wrap=True)
    if desde != hasta:
        ws.merge_cells('%s%d:%s%d' % (desde, fila, hasta, fila))
    ws.row_dimensions[fila].height = max(16, 15 * math.ceil(len(texto) / chars))


def _imprimir(ws, horizontal=False):
    ws.page_setup.paperSize = M.PAPEL_EN                                          # 1 = US Letter
    ws.page_setup.orientation = 'landscape' if horizontal else 'portrait'
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.oddFooter.center.text = FOOTER
    ws.print_options.horizontalCentered = True
    ws.protection.sheet = True                                                    # sin contraseña
    ws.sheet_view.showGridLines = False


def _dv(ws, tipo, celdas, f1=None, f2=None, op=None, titulo='Invalid value', msg=''):
    dv = DataValidation(type=tipo, formula1=f1, formula2=f2, operator=op, allow_blank=False, showErrorMessage=True,
                        errorTitle=titulo, error=msg)
    ws.add_data_validation(dv)
    for c in celdas:
        dv.add(c)


def a_excel(expr, ref):
    return re.sub(r'\{(\w+)\}', lambda m: ref[m.group(1)], expr)


def _oz_ext():
    return "'%s'!$B$%d" % (HP, M_P['OZ'])


def _bloque(ws, titulo_in, ent, m, ra, sal, s, neg, chars=105.0):
    """Escribe el bloque de entradas (verdes/grises) y el de resultados de una hoja de dos bloques."""
    _f(ws, 'A5', titulo_in, b=True, fill=GRIS)
    _f(ws, 'B5', 'Value', b=True, fill=GRIS)
    _f(ws, 'C5', 'Notes', b=True, fill=GRIS)
    for k, rot, fmt, val, const, nota in ent:
        f = m[k]
        _f(ws, 'A%d' % f, rot)
        if const:
            _f(ws, 'B%d' % f, val, fmt=fmt, fill=GRIS)
        else:
            _f(ws, 'B%d' % f, val, fmt=fmt, fill=VERDE, locked=False)
        _nota(ws, f, nota, desde='C', hasta='I', chars=chars)
    ref = dict((k, '$B$%d' % f) for k, f in m.items())
    ref.update((k, '$B$%d' % f) for k, f in s.items())
    for col, t in zip('ABC', ('RESULTS', 'Value', 'Notes')):
        _f(ws, '%s%d' % (col, ra), t, b=True, fill=GRIS)
    for k, rot, fmt, expr, nota in sal:
        f = s[k]
        _f(ws, 'A%d' % f, rot, b=k in neg)
        _f(ws, 'B%d' % f, '=' + a_excel(expr, ref), fmt=fmt, b=k in neg, wrap=(fmt == 'General'))
        _nota(ws, f, nota, desde='C', hasta='I', chars=chars)


def _semaforo(ws, celda, ok, ko):
    ws.conditional_formatting.add(celda, FormulaRule(formula=['$%s$%s="%s"' % (celda[0], celda[1:], ok)], fill=OK))
    ws.conditional_formatting.add(celda, FormulaRule(formula=['$%s$%s="%s"' % (celda[0], celda[1:], ko)], fill=ROJO))


# ---------------------------------------------------------------------------------------------------------------------
# Hojas
# ---------------------------------------------------------------------------------------------------------------------
def hoja_pollo(wb):
    ws = wb.active
    ws.title = HP
    for col, w in (('A', 60), ('B', 15), ('C', 15), ('D', 16), ('E', 14), ('F', 14), ('G', 14), ('H', 12), ('I', 15)):
        ws.column_dimensions[col].width = w
    _f(ws, 'A1', 'CHICKEN YIELD: HOW MUCH A COOKED BIRD WEIGHS, WHAT IT COSTS AND HOW MANY BIRDS YOU NEED', b=True, sz=14)
    _f(ws, 'A2', 'Type only in the GREEN cells: everything else is calculated. Amounts before sales tax. Every default value is a kit estimate, '
                 'editable: replace it with your supplier\'s numbers and, above all, with what you measure on your own birds.', wrap=True)
    ws.merge_cells('A2:I2')
    ws.row_dimensions[2].height = 32
    _f(ws, 'A3', 'FIXED NOTICE. The cooked yield (72 %) is a kit estimate; the kit works with a range of 70 to 75 % and no manufacturer figure is '
                 'published. Weigh 5 raw and cooked birds from the same load, and adjust the yield until it matches what you measure.',
       wrap=True, b=True, fill=GRIS)
    ws.merge_cells('A3:I3')
    ws.row_dimensions[3].height = 36
    _bloque(ws, 'A. THE BIRD', ENT_P, M_P, M_RP, SAL_P, S_P, ('asado', 'rend', 'coste', 'coste_kg', 'pollos_N'))
    cel = lambda k: 'B%d' % M_P[k]
    _dv(ws, 'decimal', [cel(k) for k in ('P', 'Pk')], f1='0', op='greaterThan', msg='Type a number greater than 0.')
    _dv(ws, 'decimal', [cel(k) for k in ('S', 'Ns')], f1='0', op='greaterThanOrEqual', msg='Type a number greater than or equal to 0.')
    _dv(ws, 'decimal', [cel('r')], f1='0.3', f2='1', op='between', titulo='Invalid yield', msg='Type a value between 0.3 and 1 (0.72 = 72 %).')
    _dv(ws, 'decimal', [cel('a'), cel('hw')], f1='0', f2='1', op='between', titulo='Invalid percentage', msg='Type a percentage between 0 and 1 (0.5 = 50 %).')
    _dv(ws, 'whole', [cel('pp'), cel('Np')], f1='1', f2='1000000', op='between', titulo='Invalid number', msg='Type a whole number of 1 or more.')
    _imprimir(ws, horizontal=True)


def hoja_asador(wb):
    ws = wb.create_sheet(HA)
    for col, w in (('A', 64), ('B', 15), ('C', 15), ('D', 16), ('E', 14), ('F', 14), ('G', 14), ('H', 12), ('I', 15)):
        ws.column_dimensions[col].width = w
    _f(ws, 'A1', 'ROTISSERIE LOADS: DO YOUR SPITS REACH THE SUNDAY PEAK?', b=True, sz=14)
    _f(ws, 'A2', 'Type only in the GREEN cells: everything else is calculated. Every default value is a kit estimate, editable. The capacity counts '
                 'the full loads that fit between the first load and the deadline.', wrap=True)
    ws.merge_cells('A2:I2')
    ws.row_dimensions[2].height = 32
    _f(ws, 'A3', 'FIXED NOTICE. The cook time (80 minutes) has no published figure: no spec sheet gives the time of a 3 lb bird. Time your own unit '
                 'and your own bird size before you trust the capacity.', wrap=True, b=True, fill=GRIS)
    ws.merge_cells('A3:I3')
    ws.row_dimensions[3].height = 36
    _bloque(ws, 'A. THE ROTISSERIES AND THE SUNDAY', ENT_S, M_S, M_RS, SAL_S, S_S,
            ('carga', 'cap_h', 'tandas', 'cap_dom', 'sem', 'sem_ago', 'tarde_h', 'tarde_m', 'gas_pollo'))
    for k in ('sem', 'sem_ago'):
        _semaforo(ws, 'B%d' % S_S[k], SEM_OK, SEM_KO)
        ws.row_dimensions[S_S[k]].height = 32
    cel = lambda k: 'B%d' % M_S[k]
    _dv(ws, 'whole', [cel(k) for k in ('R', 'E', 'p', 'D')], f1='1', f2='1000000', op='between', titulo='Invalid number', msg='Type a whole number of 1 or more.')
    _dv(ws, 'decimal', [cel('t')], f1='0', op='greaterThan', msg='Type a number greater than 0.')
    _dv(ws, 'decimal', [cel('c'), cel('th')], f1='0', op='greaterThanOrEqual', msg='Type a number greater than or equal to 0.')
    _dv(ws, 'whole', [cel('btu')], f1='0', f2='100000000', op='between', titulo='Invalid number', msg='Type a whole number of 0 or more.')
    _dv(ws, 'whole', [cel('hh0'), cel('hh1')], f1='0', f2='23', op='between', titulo='Invalid hour', msg='Type a whole hour from 0 to 23.')
    _dv(ws, 'whole', [cel('mm0'), cel('mm1')], f1='0', f2='59', op='between', titulo='Invalid minutes', msg='Type whole minutes from 0 to 59.')
    _dv(ws, 'decimal', [cel('hp')], f1='0', f2='1', op='between', titulo='Invalid percentage', msg='Type a percentage between 0 and 1 (0.7 = 70 %).')
    _dv(ws, 'decimal', [cel('ag')], f1='0', f2='3', op='between', titulo='Invalid percentage', msg='Type a value between 0 and 3 (1 = +100 %).')
    _imprimir(ws, horizontal=True)


def hoja_escandallo(wb):
    ws = wb.create_sheet(HE)
    for col, w in (('A', 46), ('B', 14), ('C', 7)):
        ws.column_dimensions[col].width = w
    for col in PCOL:
        ws.column_dimensions[col].width = 15
    ws.column_dimensions[ORI].width = 46
    _f(ws, 'A1', 'PLATE COSTING: WHOLE, HALF, QUARTER, FAMILY MEAL, SANDWICH OR BOWL, SIDES AND A DRINK, IN OUNCES', b=True, sz=14)
    _f(ws, 'A2', 'Type only in the GREEN cells (prices, ounces, names, bird fractions and menu prices). Amounts before sales tax except the menu '
                 'price. Prices are a kit estimate, editable: type your distributor\'s. The cost of the chicken comes from the '
                 '"%s" sheet.' % HP, wrap=True)
    ws.merge_cells('A2:%s2' % ORI)
    ws.row_dimensions[2].height = 32
    for col, t in zip('ABD', ('GENERAL DATA', 'Value', 'Notes')):
        _f(ws, '%s4' % col, t, b=True, fill=GRIS)
    for k, rot, fmt, val, nota in ESC_IN:
        f = E[k]
        _f(ws, 'A%d' % f, rot)
        _f(ws, 'B%d' % f, val, fmt=fmt, fill=VERDE, locked=False)
        _nota(ws, f, nota, desde='D', hasta=ORI, chars=150.0)
    for col, t in (('A', 'INGREDIENT (in ounces)'), ('B', '$/lb before tax'), ('C', 'Unit')):
        _f(ws, '%s%d' % (col, E_CAB), t, b=True, fill=GRIS, wrap=True)
    _f(ws, '%s%d' % (ORI, E_CAB), 'WHERE THE PRICE COMES FROM', b=True, fill=GRIS, wrap=True)
    _f(ws, 'A%d' % E_PVP, 'Menu price ($)', b=True)
    for j, (nom, (pvp, _g, _u)) in enumerate(RECETAS.items()):
        col = PCOL[j]
        _f(ws, '%s%d' % (col, E_CAB), nom, b=True, fill=VERDE, wrap=True, locked=False)
        _f(ws, '%s%d' % (col, E_PVP), pvp, fmt=USD, fill=VERDE, locked=False, b=True)
    ws.row_dimensions[E_CAB].height = 62
    for i, (nom, p, origen) in enumerate(ING_G):
        f = E_G0 + i
        _f(ws, 'A%d' % f, nom, fill=VERDE, locked=False)
        _f(ws, 'B%d' % f, p, fmt=USD, fill=VERDE, locked=False)
        _f(ws, 'C%d' % f, 'oz')
        for j, (_n, (_pvp, gr, _ud)) in enumerate(RECETAS.items()):
            _f(ws, '%s%d' % (PCOL[j], f), float(gr.get(nom, 0)), fmt='#,##0.0;-#,##0.0;"-"', fill=VERDE, locked=False)
        _f(ws, '%s%d' % (ORI, f), origen, i=True)
    _f(ws, 'A%d' % E_UH, 'INGREDIENT (in units)', b=True, fill=GRIS)
    _f(ws, 'B%d' % E_UH, '$/ea before tax', b=True, fill=GRIS)
    _f(ws, 'C%d' % E_UH, 'Unit', b=True, fill=GRIS)
    for i, (nom, p, origen) in enumerate(ING_U):
        f = E_U0 + i
        if p is None:
            _f(ws, 'A%d' % f, nom)
            _f(ws, 'B%d' % f, "='%s'!$B$%d" % (HP, S_P['coste']), fmt=USD3)
            _f(ws, 'C%d' % f, 'bird')
        else:
            _f(ws, 'A%d' % f, nom, fill=VERDE, locked=False)
            _f(ws, 'B%d' % f, p, fmt=USD3, fill=VERDE, locked=False)
            _f(ws, 'C%d' % f, 'ea')
        for j, (_n, (_pvp, _g, ud)) in enumerate(RECETAS.items()):
            _f(ws, '%s%d' % (PCOL[j], f), float(ud.get(nom, 0)), fmt='#,##0.00;-#,##0.00;"-"', fill=VERDE, locked=False)
        _f(ws, '%s%d' % (ORI, f), origen, i=True)
    _f(ws, 'A%d' % E_R, 'RESULT PER ITEM (ounces and units in the columns above)', b=True, fill=GRIS)
    for j in range(NP):
        _f(ws, '%s%d' % (PCOL[j], E_R), '=%s%d' % (PCOL[j], E_CAB), b=True, fill=GRIS, wrap=True)
    ws.row_dimensions[E_R].height = 62
    oz = _oz_ext()
    gp = "'%s'!$B$%d" % (HP, S_P['g_entero'])
    for k, rot, fmt in RES:
        f = R[k]
        neg = k in ('c_tot', 'fc')
        _f(ws, 'A%d' % f, rot, b=neg, wrap=True)
        for j in range(NP):
            c = PCOL[j]
            rg = '%s$%d:%s$%d' % (c, E_G0, c, E_G1)
            ru = '%s$%d:%s$%d' % (c, E_U0 + 1, c, E_U1)
            ex = {
                'peso': 'SUM(%s)+%s$%d*%s' % (rg, c, E_U0, gp),
                'c_pollo': '%s$%d*$B$%d' % (c, E_U0, E_U0),
                'c_acomp': 'SUMPRODUCT(%s,$B$%d:$B$%d)/%s+SUMPRODUCT(%s,$B$%d:$B$%d)' % (rg, E_G0, E_G1, oz, ru, E_U0 + 1, E_U1),
                'c_tot': '%s%d+%s%d*(1+$B$%d)' % (c, R['c_pollo'], c, R['c_acomp'], E['merm']),
                'pvp_sin': '%s%d/(1+$B$%d)' % (c, E_PVP, E['tax']),
                'fc': '%s%d/%s%d' % (c, R['c_tot'], c, R['pvp_sin']),
                'sem': 'IF(%s%d<=$B$%d,"%s",IF(%s%d<=$B$%d+$B$%d,"%s","%s"))' % (
                    c, R['fc'], E['obj'], FC_OK, c, R['fc'], E['obj'], E['tol'], FC_VIG, FC_ALTO),
                'margen': '%s%d-%s%d' % (c, R['pvp_sin'], c, R['c_tot']),
                'pvp_sug': '%s%d/$B$%d*(1+$B$%d)' % (c, R['c_tot'], E['obj'], E['tax']),
            }[k]
            _f(ws, '%s%d' % (c, f), '=' + ex, fmt=fmt, b=neg, wrap=(k == 'sem'))
    ws.row_dimensions[R['sem']].height = 32
    rng = '%s%d:%s%d' % (PCOL[0], R['sem'], PCOL[-1], R['sem'])
    top = '%s%d' % (PCOL[0], R['sem'])
    for txt, fill in ((FC_OK, OK), (FC_VIG, AMBAR), (FC_ALTO, ROJO)):
        ws.conditional_formatting.add(rng, FormulaRule(formula=['%s="%s"' % (top, txt)], fill=fill))
    # margen por canal de todos los platos
    _f(ws, 'A%d' % E_CH, 'MARGIN BY CHANNEL (every item)', b=True, fill=GRIS)
    for j in range(NP):
        _f(ws, '%s%d' % (PCOL[j], E_CH), '=%s%d' % (PCOL[j], E_CAB), b=True, fill=GRIS, wrap=True)
    ws.row_dimensions[E_CH].height = 62
    for k, rot, fmt in CANAL:
        f = CH[k]
        _f(ws, 'A%d' % f, rot, b=k.startswith('plat') or k == 'rank', wrap=True)
        for j in range(NP):
            c = PCOL[j]
            ps = '%s%d' % (c, R['pvp_sin'])
            ex = {
                'sala_e': '%s%d' % (c, R['margen']),
                'sala_p': '%s%d/%s' % (c, CH['sala_e'], ps),
                'mostr_e': '%s%d-$B$%d' % (c, CH['sala_e'], E['env']),
                'mostr_p': '%s%d/%s' % (c, CH['mostr_e'], ps),
                'plat_pv': '%s*(1+$B$%d)' % (ps, E['rec']),
                'plat_e': '%s%d*(1-$B$%d)-%s%d-$B$%d' % (c, CH['plat_pv'], E['com'], c, R['c_tot'], E['env']),
                'plat_p': '%s%d/%s%d' % (c, CH['plat_e'], c, CH['plat_pv']),
                'rank': '1+COUNTIF($%s$%d:$%s$%d,">"&%s%d)' % (PCOL[0], CH['mostr_e'], PCOL[-1], CH['mostr_e'], c, CH['mostr_e']),
            }[k]
            _f(ws, '%s%d' % (c, f), '=' + ex, fmt=fmt, b=k.startswith('plat') or k == 'rank')
    ws.row_dimensions[CH['rank']].height = 32
    _f(ws, 'A%d' % E_FIN, 'The food cost in the table is the one of the plate served dine-in. On pickup and pre-orders the packaging of the order is '
                          'deducted; on the marketplace, the price is also raised with the markup and the commission is deducted from that price. '
                          'FIXED NOTICE: on the marketplace, push the whole chicken, the family meal and the sides that travel well; fries arrive soft. '
                          'Beer, chicha morada and other drinks are not in this costing.', i=True, wrap=True)
    ws.merge_cells('A%d:%s%d' % (E_FIN, ORI, E_FIN))
    ws.row_dimensions[E_FIN].height = 48
    # validaciones
    _dv(ws, 'decimal', ['B%d' % E[k] for k in ('tax', 'tol', 'merm', 'com', 'rec')], f1='0', f2='0.9', op='between',
        titulo='Invalid percentage', msg='Type a percentage between 0 and 0.9 (0.1 = 10 %).')
    _dv(ws, 'decimal', ['B%d' % E['obj']], f1='0.01', f2='0.9', op='between', titulo='Invalid target',
        msg='The target cannot be 0 (it is a divisor): type a value between 0.01 and 0.9.')
    _dv(ws, 'decimal', ['B%d' % E['env']] + ['B%d' % f for f in range(E_G0, E_G1 + 1)] + ['B%d' % f for f in range(E_U0 + 1, E_U1 + 1)]
        + ['%s%d' % (c, f) for c in PCOL for f in [E_PVP] + list(range(E_G0, E_G1 + 1)) + list(range(E_U0, E_U1 + 1))],
        f1='0', op='greaterThanOrEqual', titulo='Invalid amount', msg='Type a number greater than or equal to 0.')
    _imprimir(ws, horizontal=True)


def _instrucciones():
    return [
        ('INSTRUCTIONS: Rotisserie Chicken Yield & Costing Calculator', True),
        ('Three tools in one workbook: the "%s" sheet tells you how much a cooked bird weighs, what each cooked pound and each portion costs, and '
         'how many birds you need for a number of guests; the "%s" sheet compares the capacity of your rotisseries with the Sunday peak and with a '
         'holiday or peak day; the "%s" sheet calculates the cost, food cost, margin and menu price of the whole chicken, the half, the quarter, '
         'the family meal, a sandwich or bowl, three sides and a drink, with the margin dine-in, on pickup and pre-orders and on a delivery '
         'marketplace. Type only in the GREEN cells; everything else is calculated. Chicken prices are in $/lb and portions in ounces (16 oz per '
         'lb, the gray constant).' % (HP, HA, HE), False),
        ('Every default value is a kit estimate, editable: they are examples consistent with the kit\'s business plan case, not anyone\'s prices. '
         'Replace them with your supplier\'s numbers and with what you measure in your own kitchen before you set prices.', False),
        ('SHEET "%s"' % HP, True),
        ('1. How to measure your cooked yield. Weigh 5 raw birds from the same load, cook the 5 together and weigh them again when they come out, '
         'drained. Divide the cooked weight by the raw weight and use the average: that is your real cooked yield. The kit works with a range of '
         '70 to 75 %; no manufacturer figure is published.', False),
        ('2. The most useful results are the cost per whole bird and per cooked pound, which feed the costing, and the birds you have to order for '
         'a pre-order or a party. Look at the cost of the leftover you throw away too: the birds that stay unsold at close and are not reused '
         'are money lost.', False),
        ('3. Food safety. Probe the thickest part of the thigh: 165 °F is the Food Code cook temperature for poultry. In the heated display, hold '
         'at 135 °F or above, and check your county\'s rules for hot holding and for time as a public health control, which usually need a written '
         'procedure and a time mark on every load (under that control, birds are thrown out at 4 hours, never cooled). Leftover you want to reuse is cooled from 135 to 70 °F within 2 hours and to 41 °F within 6 hours '
         'in total, and reheated to 165 °F for hot holding. Local rules vary: ask your health department.', False),
        ('SHEET "%s"' % HA, True),
        ('1. Type your rotisseries, the spits of each one and the birds per spit, and the cook minutes. The cook time has no published figure: time a '
         'load from the moment you light the unit until the thigh reaches 165 °F, with the bird size you buy. Add the minutes to load, unload and clean.', False),
        ('2. Type the time of the first load and the deadline by which you want the birds cooked (the Sunday lunch peak runs until about 1:30 PM). '
         'The sheet counts the full loads that fit, the capacity and whether you cover the peak of a normal Sunday and the one of a holiday or peak day.', False),
        ('3. The latest first load is the time you must load the first birds so you do not fall short at the peak. The gas per bird uses the BTU/h rating '
         'of your units, the hours they burn and the price per therm from your gas bill: it is an estimate, so check it against your bill.', False),
        ('SHEET "%s"' % HE, True),
        ('1. Type the sales tax included in your menu prices (0 % by default), the target food cost, the watch margin, the waste on the sides, the '
         'packaging, the marketplace commission and the markup.', False),
        ('2. In the table, each row is an ingredient with its price per pound before tax. The ounces of each item go in its column. The chicken is '
         'measured in fractions of a whole bird (1, 0.5, 0.25) and takes its cost from the "%s" sheet.' % HP, False),
        ('3. For each item you see the cost, the price before tax, the food cost with a traffic light, the margin and the menu price with tax you would '
         'need to hit your target. At the bottom, the margin dine-in, on pickup and pre-orders and on the marketplace, and the rank of each item by '
         'margin. The whole bird tends to land above the target; the family meal and the sandwich or bowl balance it.', False),
        ('PRICES AND SOURCES', True),
        ('Every default price is an example: "Kit estimate (Oct 2026): replace with your distributor\'s price". The sample menu prices (whole $19.00, '
         'half + 2 sides $17.50, quarter + 2 sides $13.00, family meal $34.00) come from the kit\'s business plan case. Sales tax is not part of the '
         'menu prices in this sample (0 %); rates vary by state and city, so ask your accountant.', False),
        ('Sources: the kit\'s US rotisserie chicken research (October 2026): USDA AMS wholesale prices for whole chickens, distributor spec sheets for '
         'the rotisserie capacity and the BTU rating, and menus from US Peruvian chicken restaurants. No figure in this calculator is published data '
         'about your business: replace them with your own.', False),
    ]


def hoja_instrucciones(wb):
    ws = wb.create_sheet(HI)
    ws.column_dimensions['A'].width = 110
    f = 1
    for t, b in _instrucciones():
        _f(ws, 'A%d' % f, t, b=b, sz=13 if f == 1 else 11, wrap=True)
        if not b:
            ws.row_dimensions[f].height = max(16, 15 * math.ceil(len(t) / 120.0))
        f += 2 if b and f > 1 else 1
    f += 1
    for t in ('To edit a cell that is not green: Review → Unprotect Sheet (there is no password).',
              'More templates and kits in the catalog at ' + M.URL_TIENDA, M.FIRMA_EN, VERSION_LINEA, M.AVISO_EN):
        _f(ws, 'A%d' % f, t, wrap=True, i=t == M.AVISO_EN)
        f += 1
    _imprimir(ws)


def docprops(wb):
    pr = wb.properties
    pr.title = '%s · %s' % (TITULO, PRODUCTO)
    pr.subject = '%s · v%s' % (PRODUCTO, VERSION)
    pr.keywords = KEYWORDS
    pr.description = URL
    pr.category = 'AI Chef Pro · Digital products'
    pr.creator = pr.lastModifiedBy = 'AI Chef Pro'


def generar(destino):
    wb = openpyxl.Workbook()
    hoja_pollo(wb)
    hoja_asador(wb)
    hoja_escandallo(wb)
    hoja_instrucciones(wb)
    docprops(wb)
    os.makedirs(os.path.dirname(destino), exist_ok=True)
    wb.save(destino)
    import inject_cache as IC                                                      # noqa: E402
    nf, inj, fail = IC.inject(destino)
    print('%s: formulas=%d cache_inyectado=%d fallos_pycel=%d' % (os.path.basename(destino), nf, inj, fail))
    return fail


def valores_expuestos(ruta):
    """Lee de la CACHÉ del xlsx los 5 valores expuestos: {nombre: valor}."""
    wv = openpyxl.load_workbook(ruta, data_only=True)
    return OrderedDict((n, wv[h][c].value) for n, (h, c) in EXPUESTOS.items())


# ---------------------------------------------------------------------------------------------------------------------
# Autotest
# ---------------------------------------------------------------------------------------------------------------------
#: restos de la Taco EN / Pizza EN, euros, español, promesas y reglamentación ES (se suma datos_pol.RX_POL_NO si se puede importar)
PROHIBIDOS = [r'pizz', r'neapolit', r'\bdough\b', r'wood[- ]fired', r'mozzarella', r'\btacos?\b', r'taquer', r'trompo', r'al pastor', r'tortilla',
              r'masa harina', r'€', r'\bIVA\b', r'APPCC', r'\bkg\b|\bgramos?\b|\bkW\b|\bkWh\b', r'°C', r'gluten[- ]free', r'authentic', r'mezcal',
              r'tex-?mex', r'\bpollo asado\b', r'poultry (?:shop|store)', r'\broasters?\b', r'broaster', r'PERi-PERi', r'chicken shop',
              r'\b63 ?°|\b75 ?°', r'\bcocktails?\b|\bhappy hour\b|\btapas\b',
              r'\bfreelance', r'carnet', r'tarifa plana', r'convenio|\bSMI\b|aut[oó]nomo|Veri\*?factu', r'\bchimenea\b', r'\bespeto\b|\bespada']
CELDA = {}                                                                         # clave de entrada → (hoja, celda)
for _k, _f2 in M_P.items():
    CELDA[_k] = (HP, 'B%d' % _f2)
for _k, _f2 in M_S.items():
    CELDA[_k] = (HA, 'B%d' % _f2)
for _k, _f2 in E.items():
    CELDA[_k] = (HE, 'B%d' % _f2)
for _i in range(NG):
    CELDA['p%d' % _i] = (HE, 'B%d' % (E_G0 + _i))
for _i in range(1, NU):
    CELDA['up%d' % _i] = (HE, 'B%d' % (E_U0 + _i))
for _j in range(NP):
    CELDA['pvp%d' % _j] = (HE, '%s%d' % (PCOL[_j], E_PVP))
    for _i in range(NG):
        CELDA['g%d_%d' % (_i, _j)] = (HE, '%s%d' % (PCOL[_j], E_G0 + _i))
    for _i in range(NU):
        CELDA['u%d_%d' % (_i, _j)] = (HE, '%s%d' % (PCOL[_j], E_U0 + _i))


def celdas_salida(r):
    """(hoja, celda, valor esperado, etiqueta) de todas las salidas."""
    out = []
    for k in S_P:
        out.append((HP, 'B%d' % S_P[k], r['P'][k], 'P.' + k))
    for k in S_S:
        out.append((HA, 'B%d' % S_S[k], r['S'][k], 'S.' + k))
    for i in range(NG):
        out.append((HE, 'B%d' % (E_G0 + i), r['pg'][i], 'price%d' % i))
    for i in range(NU):
        out.append((HE, 'B%d' % (E_U0 + i), r['pu'][i], 'unit_price%d' % i))
    for j in range(NP):
        for k in list(R) + list(CH):
            out.append((HE, '%s%d' % (PCOL[j], R[k] if k in R else CH[k]), r['p%d' % j][k], '%s[%d]' % (k, j)))
    return out


def _distinto(x, v):
    if isinstance(v, str):
        return x != v
    return not isinstance(x, (int, float)) or abs(float(x) - v) > 0.01


def _checks_logica(nombre, jj, err):
    r = calcular(jj)
    p, s = r['P'], r['S']
    if abs(p['asado'] - jj['P'] * jj['r']) > 1e-9:
        err.append('%s: cooked weight ≠ raw weight × yield' % nombre)
    if abs(p['c_medio'] * jj['M2'] - p['coste']) > 0.01 or abs(p['c_cuarto'] * jj['M4'] - p['coste']) > 0.01:
        err.append('%s: half × 2 or quarter × 4 ≠ whole bird' % nombre)
    if abs(p['g_medio'] * jj['M2'] - p['asado'] * jj['OZ']) > 0.01 or abs(p['g_cuarto'] * jj['M4'] - p['asado'] * jj['OZ']) > 0.01:
        err.append('%s: portion ounces × count ≠ cooked ounces' % nombre)
    if abs(p['coste_kg'] * p['asado'] - p['coste']) > 0.01:
        err.append('%s: cost per cooked lb × cooked lb ≠ bird cost' % nombre)
    if p['pollos_N'] * jj['pp'] < jj['Np'] or (p['pollos_N'] - 1) * jj['pp'] >= jj['Np']:
        err.append('%s: birds for the guests is not the ceiling of guests ÷ guests per bird' % nombre)
    if abs(s['cap_dom'] - jj['R'] * jj['E'] * jj['p'] * s['tandas']) > 1e-9 or s['carga'] != s['carga_una'] * jj['R']:
        err.append('%s: capacity ≠ rotisseries × spits × birds × loads' % nombre)
    if s['cargas_dom'] * s['carga'] < jj['D'] or (s['cargas_dom'] - 1) * s['carga'] >= jj['D']:
        err.append('%s: loads of the Sunday are not the ceiling of D ÷ birds per load' % nombre)
    if s['min_disp'] >= 0 and not (s['tandas'] * s['ciclo'] <= s['min_disp'] + 1e-9 < (s['tandas'] + 1) * s['ciclo'] + 1e-9):
        err.append('%s: loads × cycle does not fit the minutes available' % nombre)
    if abs(s['cap_h'] * s['ciclo'] - s['carga'] * jj['MH']) > 1e-6:
        err.append('%s: capacity per hour × cycle ≠ birds per load × 60' % nombre)
    if abs(s['tarde_h'] * jj['MH'] + s['tarde_m'] - s['tarde_min']) > 1e-9:
        err.append('%s: latest first load hour and minutes do not add up' % nombre)
    if abs(s['gas_dia'] - jj['btu'] * jj['R'] * s['horas_gas'] / jj['BT'] * jj['th']) > 1e-9:
        err.append('%s: gas of the day ≠ BTU × rotisseries × hours ÷ BTU per therm × price' % nombre)
    for j in range(NP):
        q = r['p%d' % j]
        frac = jj['u0_%d' % j]
        if abs(q['c_pollo'] - frac * p['coste']) > 0.01:
            err.append('%s: chicken cost of column %d ≠ fraction × bird cost' % (nombre, j))
        if not (q['plat_e'] < q['mostr_e'] < q['sala_e'] and q['plat_p'] < q['mostr_p'] < q['sala_p']):
            err.append('%s: marketplace < pickup < dine-in margin does not hold in column %d' % (nombre, j))


def autotest(destino):
    from pycel import ExcelCompiler
    err, n = [], 0
    wv = openpyxl.load_workbook(destino, data_only=True)
    wf = openpyxl.load_workbook(destino)
    base = valores_defecto()
    j2 = dict(base, P=3.5, Pk=1.9, S=0.5, hw=0.08, r=0.70, pp=4, Np=75, Ns=3.0, a=0.2, R=1, E=8, p=6, t=85.0, c=10.0, hh0=10, mm0=15,
              hh1=14, mm1=30, D=140, hp=0.6, ag=0.5, btu=150000, th=1.5, tax=0.08, obj=0.28, tol=0.04, merm=0.05, env=0.7, com=0.20,
              rec=0.10, p0=3.1, p4=2.5, up1=0.5, up2=0.8, pvp0=20.5, pvp8=3.75, g0_3=20.0, g2_4=6.0, u0_1=0.5, u2_8=2.0)
    j3 = dict(base, P=2.6, Pk=1.5, r=0.75, hw=0.0, pp=2, Np=10, Ns=0.0, R=3, E=5, p=4, t=70.0, c=20.0, hh0=9, mm0=0, hh1=14, mm1=45, D=60,
              hp=0.9, ag=0.0, btu=120000, th=1.0, tax=0.0, obj=0.4, tol=0.1, merm=0.0, env=0.4, com=0.30, rec=0.2, p1=0.9, p6=5.0, up1=0.45,
              pvp3=36.0, g1_6=12.0, g5_7=6.0, u0_7=0.2, u1_5=2.0)
    # 1. caché y pycel del juego por defecto; pycel de los otros dos con set_value
    xl = ExcelCompiler(filename=destino)
    r0 = calcular(base)
    for hoja, c, v, lab in celdas_salida(r0):
        n += 1
        for et, x in (('pycel', xl.evaluate("'%s'!%s" % (hoja, c))), ('cache', wv[hoja][c].value)):
            if _distinto(x, v):
                err.append('default · %s!%s (%s): %s %r ≠ expected %r' % (hoja, c, lab, et, x, v))
    for nombre, jj in (('set 2', j2), ('set 3', j3)):
        r = calcular(jj)
        xl2 = ExcelCompiler(filename=destino)
        for hoja, c, v, lab in celdas_salida(r0):
            xl2.evaluate("'%s'!%s" % (hoja, c))
        for k, v in jj.items():
            if v != base[k]:
                h, c = CELDA[k]
                xl2.set_value("'%s'!%s" % (h, c), v)
        for hoja, c, v, lab in celdas_salida(r):
            n += 1
            x = xl2.evaluate("'%s'!%s" % (hoja, c))
            if _distinto(x, v):
                err.append('%s · %s!%s (%s): pycel %r ≠ expected %r' % (nombre, hoja, c, lab, x, v))
    # 2. identidades de la lógica en los tres juegos
    for nombre, jj in (('default', base), ('set 2', j2), ('set 3', j3)):
        n += 12
        _checks_logica(nombre, jj, err)
    # identidades sobre la hoja viva (pycel, valores por defecto)
    ev = lambda h, c: xl.evaluate("'%s'!%s" % (h, c))
    asado, pb, rr = ev(HP, 'B%d' % S_P['asado']), ev(HP, 'B%d' % M_P['P']), ev(HP, 'B%d' % M_P['r'])
    gm, gc, ge = (ev(HP, 'B%d' % S_P[k]) for k in ('g_medio', 'g_cuarto', 'g_entero'))
    cap, ca, ta = (ev(HA, 'B%d' % S_S[k]) for k in ('cap_dom', 'carga', 'tandas'))
    n += 3
    if abs(asado - pb * rr) > 1e-6 or abs(gm * 2 - ge) > 0.01 or abs(gc * 4 - ge) > 0.01:
        err.append('Chicken Yield sheet: cooked ≠ raw × yield or half × 2 / quarter × 4 ≠ whole')
    if abs(cap - ca * ta) > 1e-9 or ca != ev(HA, 'B%d' % M_S['R']) * ev(HA, 'B%d' % M_S['E']) * ev(HA, 'B%d' % M_S['p']):
        err.append('Rotisserie sheet: capacity ≠ rotisseries × spits × birds per spit × loads')
    # 3. valores con sentido
    ent, fam = r0['p%d' % J_ENTERO], r0['p%d' % J_FAMILIA]
    if not 0.28 <= ent['fc'] <= 0.38:
        err.append('whole bird food cost %.1f %% outside 28-38 %%' % (ent['fc'] * 100))
    if abs(fam['fc'] - 0.296) > 0.003:
        err.append('family meal food cost %.2f %% ≠ about 29.6 %%' % (fam['fc'] * 100))
    if abs(r0['P']['coste'] - 5.985) > 1e-6 or abs(r0['P']['asado'] - 2.16) > 1e-6 or abs(r0['P']['coste_kg'] - 2.7708) > 0.001:
        err.append('cost %.4f / cooked %.4f / $ per cooked lb %.4f outside 5.985 / 2.16 / 2.771' % (r0['P']['coste'], r0['P']['asado'], r0['P']['coste_kg']))
    if r0['S']['carga'] != 96 or r0['S']['carga_una'] != 48 or r0['S']['tandas'] != 2 or r0['S']['cap_dom'] != 192 \
            or r0['S']['cargas_dom'] != 1 or r0['S']['cargas_ago'] != 2 or r0['S']['d_ago'] != 136:
        err.append('rotisserie defaults: 96 per load · 48 each · 2 loads · 192 · 1 load a Sunday · 2 on a peak day · 136 birds (got %r)'
                   % [r0['S'][k] for k in ('carga', 'carga_una', 'tandas', 'cap_dom', 'cargas_dom', 'cargas_ago', 'd_ago')])
    if not 60 <= r0['S']['cap_h'] <= 75 or r0['S']['sem'] != SEM_OK or r0['S']['sem_ago'] != SEM_OK:
        err.append('rotisserie defaults: capacity per hour (kit: 60-75) or traffic lights unexpected')
    if not 120 <= r0['S']['d_ago'] <= 150:
        err.append('peak-day birds outside the kit range 120-150')
    if calcular(dict(base, D=400))['S']['sem'] != SEM_KO:
        err.append('the peak traffic light does not turn to "Short at the peak" with 400 birds')
    if not 0.05 <= r0['S']['gas_pollo'] <= 0.5:
        err.append('gas per bird $%.3f outside 0.05-0.50' % r0['S']['gas_pollo'])
    medio, cuarto = r0['p1'], r0['p2']
    if abs(medio['c_pollo'] * 2 - ent['c_pollo']) > 0.01 or abs(cuarto['c_pollo'] * 4 - ent['c_pollo']) > 0.01:
        err.append('half × 2 or quarter × 4 ≠ whole bird in the costing')
    if not all(0.15 <= r0['p%d' % j]['fc'] <= 0.50 for j in range(NP)):
        err.append('some default food cost outside 15-50 %%: %r' % [round(r0['p%d' % j]['fc'], 3) for j in range(NP)])
    if sorted(r0['p%d' % j]['rank'] for j in range(NP)) != list(range(1, NP + 1)):
        err.append('the margin ranks are not 1..%d' % NP)
    # 4. los 5 valores expuestos: caché frente a calcular()
    esperados = {'BIRDS_PER_LOAD': r0['S']['carga'], 'SUNDAY_CAPACITY': r0['S']['cap_dom'], 'COOKED_YIELD': r0['P']['rend'],
                 'COST_PER_BIRD': r0['P']['coste'], 'FAMILY_MEAL_FC': fam['fc']}
    expuestos = valores_expuestos(destino)
    for nombre, e_ in esperados.items():
        n += 1
        if _distinto(expuestos[nombre], e_):
            err.append('exposed %s: cache %r ≠ %r' % (nombre, expuestos[nombre], e_))
    # 5. formato, protección, DV, fórmulas, textos
    if motor.gate_formatos(wf, FICHERO):
        err.append('gate_formatos: %r' % motor.gate_formatos(wf, FICHERO)[:3])
    dvh = dict((w.title, dict((str(rng), dv) for dv in w.data_validations.dataValidation for rng in dv.sqref.ranges))
               for w in (wf[HP], wf[HA], wf[HE]))
    for hoja, k, tipo in ((HP, 'P', 'decimal'), (HP, 'r', 'decimal'), (HP, 'Np', 'whole'), (HA, 'R', 'whole'), (HA, 'E', 'whole'), (HA, 't', 'decimal'),
                          (HA, 'hh0', 'whole'), (HA, 'mm1', 'whole'), (HA, 'D', 'whole'), (HA, 'btu', 'whole'), (HE, 'obj', 'decimal'),
                          (HE, 'tax', 'decimal'), (HE, 'com', 'decimal')):
        c = CELDA[k][1]
        dv = dvh[hoja].get(c)
        if dv is None or dv.type != tipo:
            err.append('DV missing or of another type in %s!%s (expected %s)' % (hoja, c, tipo))
    textos = []
    for w in wf.worksheets:
        if w.page_setup.paperSize != 1 or not w.sheet_properties.pageSetUpPr.fitToPage or not w.protection.sheet \
                or w.oddFooter.center.text != FOOTER:
            err.append('%s: Letter + fit + protection + footer' % w.title)
        if w.protection.password:
            err.append('%s: protected with a password' % w.title)
        for row in w.iter_rows():
            for c in row:
                v = c.value
                if v == '':
                    err.append('%s!%s: empty string' % (w.title, c.coordinate))
                if isinstance(v, str) and c.data_type != 'f':
                    textos.append((w.title, c.coordinate, v))
                if isinstance(v, str) and c.data_type == 'f':
                    textos += [(w.title, c.coordinate + ' (formula)', lit) for lit in re.findall(r'"((?:[^"]|"")*)"', v)
                               if any(ch.isalpha() for ch in lit)]
                if c.data_type == 'f':
                    if wv[w.title][c.coordinate].value is None:
                        err.append('%s!%s: formula without cache' % (w.title, c.coordinate))
                    sin_refs = re.sub(r'"[^"]*"', '', re.sub(r"'[^']+'!|\$?[A-Z]{1,2}\$?\d+", '', str(c.value)))
                    for num in re.findall(r'(?<![A-Za-z])\d+\.?\d*', sin_refs):
                        if num not in ('0', '1'):
                            err.append('%s!%s: typed number %s in %s' % (w.title, c.coordinate, num, c.value))
                if c.fill.fgColor.rgb in ('00E8F5E9', 'FFE8F5E9') and c.protection.locked:
                    err.append('%s!%s: locked green cell' % (w.title, c.coordinate))
                if c.fill.fgColor.rgb in ('00E8F5E9', 'FFE8F5E9') and c.data_type == 'f':
                    err.append('%s!%s: green cell with a formula' % (w.title, c.coordinate))
                fmt = c.number_format or ''
                if '€' in fmt or '$' in fmt:
                    err.append('%s!%s: currency symbol in the format (D8)' % (w.title, c.coordinate))
        for dv_ in w.data_validations.dataValidation:
            for t in (dv_.error, dv_.errorTitle, dv_.prompt, dv_.promptTitle):
                if t:
                    textos.append((w.title, 'DV %s' % dv_.sqref, t))
    for k in ('OZ', 'M2', 'M4', 'ref'):
        c = wf[HP][CELDA[k][1]]
        if c.fill.fgColor.rgb not in ('00F2F2F2', 'FFF2F2F2') or c.protection.locked is not True:
            err.append('gray constant %s badly marked or unlocked' % k)
    for k in ('MH', 'BT'):
        c = wf[HA][CELDA[k][1]]
        if c.fill.fgColor.rgb not in ('00F2F2F2', 'FFF2F2F2') or c.protection.locked is not True:
            err.append('gray constant %s badly marked or unlocked' % k)
    p = wf.properties
    for campo in ('title', 'subject', 'keywords', 'description', 'category'):
        textos.append(('docProps', campo, getattr(p, campo) or ''))
    rx_pol = []
    try:
        import datos_pol as D
        rx_pol = [(x[0], x[2] if len(x) > 2 else None) for x in D.RX_POL_NO]
    except Exception as e:                                                         # noqa: BLE001
        print('WARNING: datos_pol.RX_POL_NO could not be imported (%s)' % e)
    n_es = n_pz = 0
    for h, c, t in textos:
        # la palabra del estilo peruano, el único español permitido, va en las keywords de docProps
        t_es = t.replace('pollo a la brasa', '') if h == 'docProps' else t
        r = [x for x in M.restos_espanol(t_es) if not x[0].startswith('oficio')]
        if r:
            n_es += 1
            err.append('%s!%s: Spanish leftovers %s' % (h, c, r[:1]))
        if M.no_latinos(t):
            err.append('%s!%s: non-Latin characters' % (h, c))
        for rx in PROHIBIDOS:
            if re.search(rx, t, re.I if not rx.startswith('(?i)') else 0):
                n_pz += 1
                err.append('%s!%s: forbidden /%s/ in %r' % (h, c, rx, t[:60]))
        for rx, exc in rx_pol:
            if re.search(rx, t) and not (exc and re.search(exc, t)):
                n_pz += 1
                err.append('%s!%s: RX_POL_NO /%s/ in %r' % (h, c, rx, t[:60]))
        if re.search(r'(?i)\bvalue\b.*\$\d|\bsave \$|total value|\bwas \$|\d+\s?% off\b', t):
            err.append('%s!%s: price anchor' % (h, c))
    todo = '\n'.join(t for _h, _c, t in textos)
    for frase in ('16 oz per lb', 'Sales tax included in your menu prices (%)', '165 °F', '135 °F', "your county's rules", SEM_OK, SEM_KO,
                  FC_OK, FC_VIG, FC_ALTO, 'Kit estimate', "Kit estimate (Oct 2026): replace with your distributor's price"):
        if frase not in todo:
            err.append('missing text %r' % frase)
    if M.FIRMA_EN not in todo or VERSION_LINEA not in todo or M.AVISO_EN not in todo:
        err.append('missing signature, version line or notice')
    if wf[HE]['A5'].value != 'Sales tax included in your menu prices (%)' or wf[HE]['B5'].value != 0:
        err.append('Costing!B5 is not the sales tax = 0 %')
    if wf.sheetnames != [HP, HA, HE, HI]:
        err.append('sheet names %r' % wf.sheetnames)
    if NG + NU < 10:
        err.append('fewer than 10 ingredients')
    if p.creator != 'AI Chef Pro' or p.description != URL or 'AI Chef Pro' not in (p.keywords or '') \
            or not (p.title or '').startswith(TITULO):
        err.append('docProps: creator=%r title=%r description=%r keywords=%r' % (p.creator, p.title, p.description, p.keywords))
    for x in err[:80]:
        print('ERROR', x)
    print('autotest rotisserie chicken yield & costing calculator: %d checks (pycel + cache at defaults; pycel in 2 more sets; identities) '
          'over %d texts · %s' % (n, len(textos), 'GREEN' if not err else 'RED (%d errors)' % len(err)))
    P_, S_ = r0['P'], r0['S']
    print('  bird: %.2f lb → %.3f lb cooked (yield %.1f %%) · cost $%.4f · $%.3f per cooked lb · half %.1f oz $%.3f · quarter %.1f oz $%.3f'
          % (base['P'], P_['asado'], P_['rend'] * 100, P_['coste'], P_['coste_kg'], P_['g_medio'], P_['c_medio'], P_['g_cuarto'], P_['c_cuarto']))
    print('  rotisserie: %d per load (%d each) · %.1f per hour · %d loads before the deadline · capacity %d · peak %.1f · %s · peak day %d birds %s · '
          'latest first load %02d:%02d · gas $%.3f per bird' % (S_['carga'], S_['carga_una'], S_['cap_h'], S_['tandas'], S_['cap_dom'], S_['dem_punta'],
                                                              S_['sem'], S_['d_ago'], S_['sem_ago'], S_['tarde_h'], S_['tarde_m'], S_['gas_pollo']))
    for nombre, (h, c) in EXPUESTOS.items():
        print('  %-18s %s!%-5s %s' % (nombre, h, c, expuestos[nombre]))
    for j, nom in enumerate(RECETAS):
        q = r0['p%d' % j]
        print('  %-44s cost $%.3f · food cost %5.1f %% · %-14s · margin dine-in %.2f / pickup %.2f / marketplace %.2f · rank %d'
              % (nom, q['c_tot'], q['fc'] * 100, q['sem'], q['sala_e'], q['mostr_e'], q['plat_e'], q['rank']))
    return 1 if err else 0


if __name__ == '__main__':
    ap = argparse.ArgumentParser()
    ap.add_argument('--salida', required=True, help='output folder')
    ap.add_argument('--autotest', action='store_true')
    a = ap.parse_args()
    destino = os.path.join(a.salida, FICHERO)
    if not a.autotest and generar(destino):
        sys.exit('ERROR: fallos de pycel en inject_cache')
    sys.exit(autotest(destino))
