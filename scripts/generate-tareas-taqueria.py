#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""generate-tareas-taqueria.py — generador de «Tareas Recurrentes: Taquería Mexicana» (kit nº 20).

CONTRATO: las helpers de este fichero emiten el **molde v2.0 ya en estado POST-MOTOR**. Es decir,
`kit-tareas-v2_0/main.py --producto kit-tareas-taqueria --dry-run` tiene que dejar **0 diferencias**
entre lo que escribe este script y lo que el motor deja después (gate G11 de la SPEC §5).
Todo lo que el motor escribiría, lo escribe ya el generador.

Fuentes (en este orden de autoridad):
  1. `scripts/productos-digitales/kit-tareas-taqueria/03-contrato-molde-v2.md`  (manda)
  2. `scripts/productos-digitales/kit-tareas-taqueria/molde-referencia-base.json` + los xlsx de
     `astro-site/public/dl/kit-tareas/` (el kit BASE v2.0, NO el sushi-bar: ése es v1.1 de CB)
  3. `scripts/productos-digitales/kit-tareas-v2_0/motor.py` (lo que el motor ESPERA y REESCRIBE)
  4. `scripts/productos-digitales/kit-tareas-taqueria/02-SPEC-kit-tareas-taqueria.md` §3 (nombres)

Uso:
    /usr/local/bin/python3 scripts/generate-tareas-taqueria.py                  # los 11 a public/dl/
    /usr/local/bin/python3 scripts/generate-tareas-taqueria.py --dir /tmp/x     # a otra carpeta
    /usr/local/bin/python3 scripts/generate-tareas-taqueria.py --prueba --dir /tmp/x   # sólo el 01

CONTENIDO: las ~300 tareas las escriben los redactores de F2 rellenando `FICHEROS` (más abajo).
Lo que hay ahora son EJEMPLOS mínimos, marcados `# EJEMPLO — sustituir en F2`.

Cómo verificar (SIN tocar `astro-site/public/dl/`; `--origen` del orquestador lo permite):
    S=<scratchpad>
    /usr/local/bin/python3 scripts/generate-tareas-taqueria.py --dir $S/kit/kit-tareas-taqueria
    /usr/local/bin/python3 scripts/productos-digitales/kit-tareas-taqueria/comparar_molde.py \
        $S/kit/kit-tareas-taqueria/01-apertura-cierre-taqueria.xlsx
    CLAUDE_SCRATCHPAD=$S /usr/local/bin/python3 \
        scripts/productos-digitales/kit-tareas-v2_0/main.py --producto kit-tareas-taqueria \
        --dry-run --origen $S/kit/kit-tareas-taqueria
    # y el diff generador ↔ post-motor, que es lo que el contrato llama «0 diferencias»:
    #   main.diff_digest(main.digest(<origen>/f), main.digest($S/dryrun-kt/kit-tareas-taqueria/f), f)

═══ LO QUE UN REDACTOR TIENE QUE SABER (trampas medidas sobre motor.py) ═══
  · Unidades: `63 °C` con espacio NORMAL U+0020 · menos U+2212 (`−18 °C`) · `Nº` con U+00BA.
    NUNCA U+202F (espacio fino) ni U+2011 (guion no separable): el motor los reescribe y el gate
    de idempotencia se pone rojo (D14/D15 de la SPEC; `motor.py:237` `RX_GRADOS`).
  · Cualquier celda que diga «APPCC» recibe del motor la coletilla `SUFIJO_APPCC` (`motor.py:2419`).
    Este generador la añade solo (`_estable()`), así que escribe el texto natural y no la pongas.
  · Una tarea de temperatura con verbo de registro + equipo de frío recibe del motor
    «(refrigeración 0-4 °C) — anota la lectura: ____ °C» (`motor.py:2429` `texto_temperatura`).
    Eso NO se replica aquí: escribe la tarea YA en su forma final, con la cola `— anota la lectura:
    ____ °C`, o el dry-run la cambiará.
  · La columna E se titula por su CONTENIDO (`motor.py:2261` `cadencia`): todo horas → «Hora Límite»;
    días de la semana → «Día»; «mensual/trimestral…» → «Cadencia»; «antes de…» → «Antelación»;
    mezcla → «Cuándo». Pasa `col_tiempo=` acorde Y `fila2=` acorde (`FILA2_*`), porque el motor
    reescribe también la fila 2 (`motor.py:2320`, `SUBTITULO`).
  · El bloque «Se conecta con» de `Instrucciones` lo compone el motor leyendo el kit entero
    (`motor.py:3815` `_bloque_conecta`). Aquí se emite ya compuesto; si cambian los nombres de
    fichero o la cadencia de las hojas de 07/08, hay que RE-MEDIRLO con un dry-run (ver `conecta()`).
"""
import argparse
import math
import os
import sys
import re
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Protection, Side
from openpyxl.formatting.rule import FormulaRule
from openpyxl.worksheet.page import PageMargins
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

# ==========================================================================
# Identidad del producto
# ==========================================================================
PID = 'kit-tareas-taqueria'
KIT = 'Taquería Mexicana'
RAIZ = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUTPUT_DIR = os.path.join(RAIZ, 'astro-site', 'public', 'dl', PID)

#: `motor.contexto` saca el nombre del kit del TÍTULO del documento con el patrón
#: «<nombre> — <kit> · <sufijo>» (`motor.py:1604-1612`), y `set_metadata` compone
#: `subject = f'{sufijo} · v2.0'` (`motor.py:4471`). De ahí salen a la vez el pie y el subject
#: que pide el contrato §1, sin escribir ninguno a mano.
SUFIJO = 'Kit de Tareas Recurrentes: ' + KIT
SUBJECT = SUFIJO + ' · v2.0'
CREATOR = 'AI Chef Pro'
#: `keywords_ok` sólo exige que termine en «AI Chef Pro» (`motor.py:4429`): así el motor no pisa
#: unas keywords escritas a mano, que es lo que pasó en `kit-tareas-pasteleria`.
KEYWORDS = 'taquería, trompo, tortillería, salsas, checklist, tareas, AI Chef Pro'

PIE = '— Kit de Tareas Recurrentes · ' + KIT + ' · AI Chef Pro · aichef.pro'
BIO = ('Diseñado por John Guerrero — chef y consultor gastronómico desde 2010, '
       'en cocina desde los 17 años · johnguerrero.es')
CONTACTO = 'Contacto: info@aichef.pro'

#: ⚠️ CONFLICTO CONOCIDO. El contrato §2 pide «Versión 2.0 · septiembre 2026 · …», pero
#: `motor.version_line()` (`motor.py:4246`) escribe el mes HARDCODEADO en «agosto». El motor
#: RECONSTRUYE la hoja «Instrucciones» entera, así que lo que emita el generador se pisa: con
#: «septiembre» el gate de idempotencia daría 11 diferencias (una por fichero) sin que haya nada
#: roto. Se emite lo que emite el motor y se deja documentado: si se quiere «septiembre», lo que
#: hay que tocar es `motor.version_line()`, no este fichero.
VERSION = 'Versión 2.0 · septiembre 2026 · aichef.pro/' + PID + ' · info@aichef.pro'
MES_MOTOR = 'agosto'                       # `motor.version_line()`, literal
VERSION_MOTOR = 'Versión 2.0 · %s 2026 · aichef.pro/%s · info@aichef.pro' % (MES_MOTOR, PID)

# ==========================================================================
# Paleta y tipografía del molde (contrato §2/§3; medidos en kit-tareas/01)
# ==========================================================================
ORO = '00FFD700'
GRIS_SUB = '00888888'
GRIS_PIE = '00999999'
GRIS_NUM = '00666666'
BLANCO = '00FFFFFF'
NEGRO_CAB = '2D2D2D'        # fill de la fila de cabecera
VERDE = 'E8F5E9'            # celda editable
VERDE_OK = 'C8E6C9'         # formato condicional de fila completada
SECCION = 'E3F2FD'          # banda de sección
GRIS = 'F5F5F5'             # rayado del calendario
FUENTE = 'Calibri'

DV_LISTA = '"✓,—,N/A"'
DV_ERROR_TIT = 'Marca no válida'
DV_ERROR = ('Usa el desplegable: ✓, — o N/A. N/A = no aplica, sale del '
            'total; — = no hecha, cuenta como pendiente.')
HOLGURA = 5                 # filas libres DENTRO del rango del contador
ETIQ_CONTADOR = 'Tareas completadas:'
CAB_CHECKLIST = ('Nº', 'Tarea', 'Zona', 'Responsable', 'Hora Límite',
                 '✓ Completada', 'Firma')
ANCHOS_CHECKLIST = (5, 45, 12, 18, 12, 12, 15)
ANCHOS_BRIEFING = (5, 30, 50)
ANCHOS_CALENDARIO = (15, 25, 40, 20)
CAB_CALENDARIO = ('Mes', 'Fecha / Evento', 'Tareas Clave', 'Antelación')
ANCHO_B_INSTR = 100

BORDE = Border(*[Side(style='thin', color='00D0D0D0')] * 4)
PIE_PAGINA = 'AI Chef Pro · aichef.pro · Página &P de &N'
MARGENES = dict(left=0.59, right=0.59, top=0.59, bottom=0.59, header=0.3, footer=0.3)

# --- fuentes de «Instrucciones» (motor.py:3633-3636) ---------------------
F_TITULO = Font(name=FUENTE, bold=True, size=18, color=ORO)
F_CAB = Font(name=FUENTE, bold=True, size=12, color='00333333')
F_TXT = Font(name=FUENTE, size=11, color='00555555')

# --- filas 2 canónicas ----------------------------------------------------
#: Las tres últimas son los literales EXACTOS de `motor.SUBTITULO` (`motor.py:2238`): si la
#: columna E lleva días / cadencias / antelaciones, el motor reescribe la fila 2 con ellos.
FILA2_TURNO = ('Fecha: ___/___/______    Turno: ☐ Almuerzo  ☐ Cena    '
               'Responsable: _________________________')
FILA2_ANUAL = 'Año: ______    Taquería: _________________________'
FILA2_DIA = ('Semana del ___/___/______ al ___/___/______    '
             'Responsable: _________________________')
FILA2_CADENCIA = ('Mes: ________________   Año: ______    '
                  'Responsable: _________________________')
FILA2_ANTELACION = ('Evento / temporada: ___________________    '
                    'Fecha: ___/___/______    '
                    'Responsable: _________________________')

# ==========================================================================
# Estabilidad de texto — lo que el motor normalizaría, ya normalizado
# ==========================================================================
#: Copias LITERALES de `motor.py:237-247` y `motor.py:2418-2426`. Si el motor cambia, estas
#: cuatro líneas se re-copian: son la diferencia entre 0 y N diferencias en el gate G11.
RX_GRADOS = re.compile(r'(?<=\d)\s*°\s*C\b')
RX_MENOS = re.compile(r'(?<![\d\w])-(?=\s?\d+(?:[.,]\d+)?\s?°C)')
RX_MENOS_RANGO = re.compile(
    r'(?<![\d\w])-(?=\s?\d+(?:[.,]\d+)?\s*(?:a|y|hasta|hacia)\s*'
    r'[−-]?\s?\d+(?:[.,]\d+)?\s?°C)')
RX_APPCC = re.compile(r'APPCC')
SUFIJO_APPCC = (' (si tienes el Pack APPCC, regístralo allí; si no, anótalo '
                'junto a la tarea)')


def _estable(v):
    """El texto tal y como quedará TRAS el motor (`motor.forma_estable`, sin temperaturas).

    `texto_temperatura` NO se replica a propósito: depende de tres regex acopladas y de la
    columna, y su salida es contenido editorial («— anota la lectura: ____ °C») que el redactor
    tiene que ver escrito. Si una tarea la dispara, el dry-run lo canta como diferencia.
    """
    if not isinstance(v, str):
        return v
    if '°' in v:
        v = RX_MENOS_RANGO.sub('−', RX_MENOS.sub('−', RX_GRADOS.sub(' °C', v)))
    if RX_APPCC.search(v) and 'si tienes el Pack APPCC' not in v \
            and 'Si tienes el Pack APPCC' not in v:
        v = v + SUFIJO_APPCC
    return v


# ==========================================================================
# Helpers de bajo nivel
# ==========================================================================
def _fill(color):
    return PatternFill('solid', fgColor=color)


def _celda(ws, fila, col, valor=None, size=11, bold=False, color=None,
           fill=None, halign=None, valign=None, wrap=False, borde=False):
    c = ws.cell(row=fila, column=col)
    if valor is not None:
        c.value = _estable(valor)
    c.font = Font(name=FUENTE, size=size, bold=bold, color=color)
    if fill:
        c.fill = _fill(fill)
    if halign or valign or wrap:
        c.alignment = Alignment(horizontal=halign, vertical=valign, wrap_text=wrap)
    if borde:
        c.border = BORDE
    return c


def _anchos(ws, anchos):
    for i, w in enumerate(anchos, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def _print_setup(ws, header_row=None, landscape=False):
    """Calco de `motor.print_setup` (`motor.py:539`)."""
    ws.page_setup.paperSize = 9                       # A4
    ws.page_setup.orientation = 'landscape' if landscape else 'portrait'
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_margins = PageMargins(**MARGENES)
    ws.oddFooter.center.text = PIE_PAGINA
    ws.oddFooter.center.size = 8
    if header_row:
        ws.print_title_rows = '%d:%d' % (header_row, header_row)
        ws.freeze_panes = ws.cell(row=header_row + 1, column=1).coordinate


def _area_impresion(ws):
    """Calco de `motor.area_impresion`: hasta la última celda CON VALOR."""
    max_r = max_c = 0
    for row in ws.iter_rows():
        for c in row:
            if c.value is not None:
                max_r, max_c = max(max_r, c.row), max(max_c, c.column)
    if max_r:
        ws.print_area = 'A1:%s%d' % (get_column_letter(max_c), max_r)


def _proteger(ws, cuerpo):
    """Calco de `motor.proteger` (`motor.py:4361`): cuerpo y verdes desbloqueados, sin contraseña."""
    for row in ws.iter_rows():
        for c in row:
            dentro = bool(cuerpo) and cuerpo[0] <= c.row <= cuerpo[1]
            verde = (c.fill is not None and c.fill.fill_type == 'solid'
                     and isinstance(c.fill.fgColor.rgb, str)
                     and c.fill.fgColor.rgb.upper().endswith(VERDE))
            c.protection = Protection(locked=not (dentro or verde))
    ws.protection.sheet = True
    ws.protection.selectLockedCells = False
    ws.protection.selectUnlockedCells = False
    ws.protection.formatCells = False
    ws.protection.formatColumns = False
    ws.protection.formatRows = False
    ws.protection.insertRows = False          # False = PERMITIDO insertar
    ws.protection.deleteRows = False
    ws.protection.sort = False
    ws.protection.autoFilter = False


# ==========================================================================
# Bloques de «Instrucciones» que EMITE EL MOTOR — se copian literales
# ==========================================================================
#: `motor._bloques_contador()` (`motor.py:3767`). Si no salieran ya escritos, el motor los
#: añadiría en la 1.ª pasada y la comparación 1.ª↔2.ª seguiría verde, pero el diff contra el
#: fichero generado no: el contrato exige el fichero YA en estado post-motor.
BLOQUES_CONTADOR = [
    ('h', 'Cómo cuenta el contador'),
    ('b', 'Marca con ✓ en la columna «✓ Completada» (desplegable): es la '
          'que cuenta el total de tareas completadas.'),
    ('b', '«N/A» = no aplica en tu local: esa tarea SALE del total (no la '
          'borres, márcala N/A). «—» = no hecha: sigue contando como '
          'pendiente y baja el porcentaje, que es de lo que se trata.'),
    ('b', 'El denominador cuenta las tareas escritas en la columna '
          '«Tarea», no un número fijo: si añades o borras tareas, el '
          'total se ajusta solo.'),
    ('h', 'Filas libres'),
    ('b', 'Al final de cada tabla hay %d filas verdes libres '
          'DENTRO del rango que cuenta el contador: escribe ahí tus '
          'tareas propias.' % HOLGURA),
    ('b', 'Si necesitas más, inserta filas DENTRO de la tabla (clic '
          'derecho → Insertar), nunca por debajo de la última: lo que se '
          'escriba fuera del rango no lo cuenta nadie.'),
]

#: `motor._bloque_proteccion()` (`motor.py:3804`).
BLOQUE_PROTECCION = [
    ('h', 'Protección de la hoja'),
    ('b', 'Las hojas con fórmulas van protegidas SIN contraseña: las '
          'celdas de entrada (las verdes) están desbloqueadas y los '
          'cálculos no se pisan por accidente.'),
    ('b', 'Puedes insertar y borrar filas con la protección puesta. Para '
          'cambiar la estructura: Revisar → Desproteger hoja.'),
]

#: Encabezados que `motor.MIS_BLOQUES` (`motor.py:3644`) DESCARTA al releer la hoja. Un bloque
#: propio no puede llamarse así o desaparecería en la 1.ª pasada sin dejar rastro.
PROHIBIDOS = ('Cómo cuenta el contador', 'Filas libres', 'Protección de la hoja',
              'Se conecta con', 'Dónde encaja este fichero', 'Qué resuelve',
              'Cómo usar', 'Celdas editables', 'Registro Mensual')
#: `motor.RX_OBSOLETO` (`motor.py:3760`): viñetas que el motor borra por contradecir a los
#: bloques nuevos. No escribirlas.
RX_OBSOLETO = re.compile(
    r'(?i)(borra las tareas que no aplican|a[ñn]ade tareas espec[ií]ficas'
    r'|filas vac[ií]as del final|hojas en blanco)')


# ==========================================================================
# HELPER 1 · hoja «Instrucciones»
# ==========================================================================
def hoja_instrucciones(wb, titulo, bloques_extra, se_conecta_con,
                       con_checklist=True):
    """Primera hoja de los 11 ficheros, con el molde ▸ del kit base v2.0.

    `bloques_extra`  : [('h', 'Encabezado'), ('b', 'viñeta'), …] — los bloques PROPIOS del
                       fichero (orden del contrato §2: «Cómo usar estas plantillas», «Cómo
                       personalizar», glosario…). Sin el «▸»: lo pone el pintado.
    `se_conecta_con` : ('Se conecta con' | 'Dónde encaja este fichero', [viñetas]) — lo que el
                       motor compone en `_bloque_conecta`. Ver `conecta()`.
    `con_checklist`  : el fichero tiene alguna hoja de checklist → se emiten los bloques del
                       contador, de las filas libres y de la protección (`motor.py:4189-4191`).

    El pintado es un calco de `motor.reescribir_instrucciones` (`motor.py:4108`), incluidas las
    alturas de fila: `main.digest` las compara una a una.
    """
    ws = wb['Instrucciones'] if 'Instrucciones' in wb.sheetnames \
        else wb.create_sheet('Instrucciones', 0)
    for tipo, texto in bloques_extra:
        if tipo == 'h' and texto.strip().rstrip(':') in PROHIBIDOS:
            raise SystemExit('Encabezado reservado del motor: «%s» (se borraría '
                             'en el dry-run). Ver PROHIBIDOS.' % texto)
        if tipo == 'b' and RX_OBSOLETO.search(texto):
            raise SystemExit('Viñeta que el motor borra por obsoleta: «%s».' % texto)

    bloques = list(bloques_extra)
    if con_checklist:
        bloques += BLOQUES_CONTADOR + BLOQUE_PROTECCION
    rotulo, vinetas = se_conecta_con
    bloques += [('h', rotulo)] + [('b', v) for v in vinetas]

    ws.column_dimensions['A'].width = 3
    ws.column_dimensions['B'].width = ANCHO_B_INSTR

    def escribe(fila, texto, fuente):
        texto = _estable(texto)
        cel = ws.cell(row=fila, column=2, value=texto)
        cel.font = fuente
        cel.alignment = Alignment(wrap_text=True, vertical='top')
        size = fuente.size or 11
        cap = max(20, int(ANCHO_B_INSTR * 11 / size))
        alto = 15 * size / 11.0
        ws.row_dimensions[fila].height = max(
            15, math.ceil(math.ceil(len(texto) / cap) * alto * 1.05))

    escribe(2, titulo, F_TITULO)
    ws.row_dimensions[2].height = 26
    fila = 4
    for tipo, texto in bloques:
        if tipo == 'h':
            if fila > 4:
                fila += 1
            escribe(fila, texto, F_CAB)
            fila += 2
        else:
            escribe(fila, '▸ ' + texto, F_TXT)
            fila += 1
    fila += 1
    escribe(fila, PIE, F_CAB)
    fila += 2
    escribe(fila, BIO, F_CAB)
    fila += 1
    escribe(fila, CONTACTO, F_CAB)
    fila += 2
    escribe(fila, VERSION_MOTOR, F_CAB)
    ws.page_setup.paperSize = 9
    ws.page_setup.orientation = 'portrait'
    ws.print_area = 'A1:B%d' % fila
    return ws


# ==========================================================================
# HELPER 2 · hoja de checklist (molde P4 del contrato §3)
# ==========================================================================
def _rejilla_checklist(ws, titulo_fila1, fila2, col_tiempo):
    _anchos(ws, ANCHOS_CHECKLIST)
    _celda(ws, 1, 1, titulo_fila1, size=16, bold=True, color=ORO)
    ws.merge_cells('A1:G1')
    _celda(ws, 2, 1, fila2, size=11, color=GRIS_SUB)
    ws.merge_cells('A2:G2')
    cab = list(CAB_CHECKLIST)
    cab[4] = col_tiempo
    for i, texto in enumerate(cab, start=1):
        _celda(ws, 4, i, texto, size=11, bold=True, color=BLANCO,
               fill=NEGRO_CAB, halign='center', valign='center', wrap=True,
               borde=True)


def _fila_seccion(ws, fila, nombre):
    _celda(ws, fila, 1, nombre, size=12, bold=True, color=ORO, fill=SECCION,
           borde=True)
    for c in range(2, 8):
        _celda(ws, fila, c, borde=True)
    ws.merge_cells('A%d:G%d' % (fila, fila))


def _fila_tarea(ws, fila, n, tarea, zona, responsable, hora):
    _celda(ws, fila, 1, n, size=10, color=GRIS_NUM, halign='center',
           valign='center', borde=True)
    _celda(ws, fila, 2, tarea, size=11, halign='left', valign='center',
           wrap=True, borde=True)
    _celda(ws, fila, 3, zona, size=10, color=GRIS_NUM, fill=VERDE,
           halign='center', valign='center', wrap=True, borde=True)
    _celda(ws, fila, 4, responsable, size=11, fill=VERDE, halign='center',
           valign='center', wrap=True, borde=True)
    _celda(ws, fila, 5, hora, size=11, fill=VERDE, halign='center',
           valign='center', wrap=True, borde=True)
    _celda(ws, fila, 6, None, size=11, fill=VERDE, halign='center',
           valign='center', wrap=True, borde=True)
    _celda(ws, fila, 7, None, size=11, fill=VERDE, borde=True)


def _fila_libre(ws, fila, con_tarea_verde=False):
    """Las 5 filas libres: mismo estilo que una fila de tarea + TODAS las columnas verdes.

    Es lo que hace `motor.normalizar_checklist` (`motor.py:1996-2007`): copia el estilo de la
    última fila numerada y `_verde()` las siete columnas. SIN número en A — si lo llevaran,
    `geometria` las contaría como tareas y cada pasada añadiría 5 más.
    """
    _celda(ws, fila, 1, None, size=10, color=GRIS_NUM, fill=VERDE,
           halign='center', valign='center', borde=True)
    _celda(ws, fila, 2, None, size=11, fill=VERDE, halign='left',
           valign='center', wrap=True, borde=True)
    _celda(ws, fila, 3, None, size=10, color=GRIS_NUM, fill=VERDE,
           halign='center', valign='center', wrap=True, borde=True)
    for c in (4, 5, 6):
        _celda(ws, fila, c, None, size=11, fill=VERDE, halign='center',
               valign='center', wrap=True, borde=True)
    _celda(ws, fila, 7, None, size=11, fill=VERDE, borde=True)
    if con_tarea_verde:
        pass                                   # B ya va verde en las libres


def _cola_checklist(ws, hr, fin):
    """Contador honesto + «Verificado por / Firma» + pie. `motor._contador` (`motor.py:2180`)."""
    contador = fin + 2
    tarea, marca = 'B', 'F'
    num = '=COUNTIFS(%s%d:%s%d,"?*",%s%d:%s%d,"✓")' % (
        tarea, hr + 1, tarea, fin, marca, hr + 1, marca, fin)
    den = '=COUNTIF(%s%d:%s%d,"?*")-COUNTIF(%s%d:%s%d,"N/A")' % (
        tarea, hr + 1, tarea, fin, marca, hr + 1, marca, fin)
    _celda(ws, contador, 1, ETIQ_CONTADOR, size=11, bold=True)
    ws.merge_cells('A%d:C%d' % (contador, contador))
    _celda(ws, contador, 4, num, size=11, bold=True)
    _celda(ws, contador, 5, 'de', size=11)
    _celda(ws, contador, 6, den, size=11, bold=True)

    verif = contador + 2
    _celda(ws, verif, 1, 'Verificado por:', size=11, bold=True)
    _celda(ws, verif, 2, None, size=11, fill=VERDE, borde=True)
    ws.merge_cells('B%d:D%d' % (verif, verif))
    _celda(ws, verif, 5, 'Firma:', size=11, bold=True)
    _celda(ws, verif, 6, None, size=11, fill=VERDE, borde=True)
    ws.merge_cells('F%d:G%d' % (verif, verif))

    pie = verif + 2
    _celda(ws, pie, 1, PIE, size=9, color=GRIS_PIE)
    return pie


def _dv_y_cf(ws, filas_dv, hr, fin):
    dv = DataValidation(type='list', formula1=DV_LISTA, allow_blank=True,
                        showErrorMessage=True, errorStyle='stop',
                        errorTitle=DV_ERROR_TIT, error=DV_ERROR)
    ws.add_data_validation(dv)
    for r in filas_dv:
        dv.add(ws.cell(row=r, column=6))
    ws.conditional_formatting.add(
        'A%d:G%d' % (hr + 1, fin),
        FormulaRule(formula=['$F%d="✓"' % (hr + 1)],
                    fill=PatternFill('solid', start_color=VERDE_OK,
                                     end_color=VERDE_OK)))


def hoja_checklist(wb, nombre, fila2, secciones, col_tiempo='Hora Límite',
                   titulo=None):
    """Hoja de checklist del molde v2.0, YA en estado post-motor.

    `secciones` = [(NOMBRE_DE_SECCION, [(tarea, zona, responsable, hora), …]), …]
       · el NOMBRE va SIN los dos espacios de sangría: los pone la helper.
       · la numeración de A es CONTINUA por hoja (no se reinicia en cada sección).
    `col_tiempo` = rótulo de la columna E. Tiene que casar con el CONTENIDO (`motor.cadencia`).
    `titulo`     = fila 1; por defecto «Checklist: <nombre>» (SPEC §3, 06).
    """
    ws = wb.create_sheet(nombre)
    _rejilla_checklist(ws, titulo or ('Checklist: ' + nombre), fila2, col_tiempo)
    hr, fila, n = 4, 5, 0
    filas_dv = []
    for seccion, tareas in secciones:
        _fila_seccion(ws, fila, '  ' + seccion)
        fila += 1
        for tarea, zona, responsable, hora in tareas:
            n += 1
            _fila_tarea(ws, fila, n, tarea, zona, responsable, hora)
            filas_dv.append(fila)
            fila += 1
    if not n:
        raise SystemExit('«%s» sin ninguna fila numerada: `motor.geometria` '
                         'devolvería None y la hoja quedaría fuera de alcance.' % nombre)
    for _ in range(HOLGURA):
        _fila_libre(ws, fila)
        filas_dv.append(fila)
        fila += 1
    fin = fila - 1
    _dv_y_cf(ws, filas_dv, hr, fin)
    pie = _cola_checklist(ws, hr, fin)
    _print_setup(ws, header_row=hr, landscape=False)
    _area_impresion(ws)
    _proteger(ws, (hr + 1, fin))
    return ws


# ==========================================================================
# HELPER 3 · 09-plantilla-personalizable
# ==========================================================================
def hoja_plantilla(wb, nombre, fila2, secciones, filas_por_seccion=5,
                   col_tiempo='Cuándo', titulo=None):
    """Plantilla en blanco: misma rejilla, secciones a personalizar y filas NUMERADAS vacías.

    ⚠️ Las filas van NUMERADAS (15 = 3 secciones × 5), no en blanco sin número, aunque el
    contrato §4 diga lo contrario: `motor.geometria` (`motor.py:800-808`) mide la última fila del
    cuerpo por el ENTERO de la columna A y devuelve `None` si no hay ninguno — la hoja quedaría
    fuera de alcance, sin DV, sin contador y sin protección. Es también lo que hace el molde
    medido (`kit-tareas/07-plantilla-personalizable.xlsx`: 15 numeradas + 5 libres).
    """
    ws = wb.create_sheet(nombre)
    _rejilla_checklist(ws, titulo or ('Plantilla en Blanco — ' + KIT), fila2, col_tiempo)
    hr, fila, n = 4, 5, 0
    filas_dv = []
    for seccion in secciones:
        _fila_seccion(ws, fila, '  ' + seccion)
        fila += 1
        for _ in range(filas_por_seccion):
            n += 1
            _fila_tarea(ws, fila, n, None, None, None, None)
            # en la plantilla también «Tarea» es una celda de entrada
            ws.cell(row=fila, column=2).fill = _fill(VERDE)
            filas_dv.append(fila)
            fila += 1
    for _ in range(HOLGURA):
        _fila_libre(ws, fila)
        filas_dv.append(fila)
        fila += 1
    fin = fila - 1
    _dv_y_cf(ws, filas_dv, hr, fin)
    _cola_checklist(ws, hr, fin)
    _print_setup(ws, header_row=hr, landscape=False)
    _area_impresion(ws)
    _proteger(ws, (hr + 1, fin))
    return ws


# ==========================================================================
# HELPER 4 · BONUS-01 briefing
# ==========================================================================
def hoja_briefing(wb, bloques, nombre='Briefing Pre-Servicio',
                  titulo='Briefing Pre-Servicio — ' + KIT, fila2=None,
                  firma='Firma encargado: _________________________'):
    """Formulario del briefing (contrato §5). `bloques` = [(RÓTULO, [etiquetas…]), …].

    ⚠️ El título de la fila 1 va en la COLUMNA B y tiene que contener «Briefing» con esa
    capitalización exacta: `motor.es_briefing` (`motor.py:1249`) hace `'Briefing' in b1`, que es
    sensible a mayúsculas. Con el `BRIEFING PRE-SERVICIO` en versalitas que pedía la SPEC §3 la
    hoja NO se reconocería: se publicaría sin print_area, sin protección, sin verdes y sin pie.
    `b2` tiene que empezar por «Fecha:».
    Cada etiqueta acabada en «:» recibe su celda verde en C (`motor.briefing`, `motor.py:3609`).
    """
    ws = wb.create_sheet(nombre)
    _anchos(ws, ANCHOS_BRIEFING)
    _celda(ws, 1, 2, titulo, size=16, bold=True, color=ORO)
    ws.merge_cells('B1:C1')
    _celda(ws, 2, 2, fila2 or FILA2_TURNO.replace('Responsable:', 'Responsable briefing:'),
           size=11, color=GRIS_SUB)
    fila = 4
    for rotulo, etiquetas in bloques:
        _celda(ws, fila, 2, rotulo, size=11, bold=True, color=BLANCO, fill=NEGRO_CAB)
        _celda(ws, fila, 3, None, fill=NEGRO_CAB)
        ws.merge_cells('B%d:C%d' % (fila, fila))
        fila += 1
        for etiqueta in etiquetas:
            _celda(ws, fila, 2, etiqueta, size=11, bold=True)
            _celda(ws, fila, 3, None, size=11, fill=VERDE, borde=True)
            fila += 1
        fila += 1
    fila += 1
    _celda(ws, fila, 2, firma, size=11, color=GRIS_SUB)
    pie = fila + 2
    _celda(ws, pie, 2, PIE, size=9, color=GRIS_PIE)
    _print_setup(ws, header_row=None, landscape=False)
    _area_impresion(ws)
    _proteger(ws, None)                 # sólo las verdes quedan desbloqueadas
    return ws


# ==========================================================================
# HELPER 5 · BONUS-02 calendario anual
# ==========================================================================
def hoja_calendario(wb, filas, nombre='Calendario Anual',
                    titulo='Calendario Anual — Fechas Clave de la Taquería',
                    fila2=None, libres=5):
    """Molde CALENDARIO de la familia NO-CB (contrato §6). `filas` = [(mes, evento, tareas, antelación), …].

    El rayado se repinta POR ÍNDICE (gris / blanco) porque eso es lo que hace `motor.calendario`
    (`motor.py:3592-3601`): emitirlo de otra forma daría diferencia en la 1.ª pasada.
    Las filas libres llevan los cuatro placeholders del molde, en verde.
    """
    ws = wb.create_sheet(nombre)
    _anchos(ws, ANCHOS_CALENDARIO)
    _celda(ws, 1, 1, titulo, size=16, bold=True, color=ORO)
    ws.merge_cells('A1:D1')
    _celda(ws, 2, 1, fila2 or FILA2_ANUAL, size=11, color=GRIS_SUB)
    for i, texto in enumerate(CAB_CALENDARIO, start=1):
        _celda(ws, 4, i, texto, size=11, bold=True, color=BLANCO,
               fill=NEGRO_CAB, halign='center', valign='center', wrap=True,
               borde=True)
    fila = 5
    for i, (mes, evento, tareas, antelacion) in enumerate(filas):
        color = GRIS if i % 2 == 0 else 'FFFFFF'
        for c, valor in enumerate((mes, evento, tareas, antelacion), start=1):
            _celda(ws, fila, c, valor, size=11, fill=color,
                   valign='center', wrap=(c == 3), borde=True)
        fila += 1
    primera_libre = fila
    for _ in range(libres):
        for c, valor in enumerate(('(Tu fecha)', '(Añade aquí)',
                                   '(Tareas específicas)', '(Antelación)'), start=1):
            _celda(ws, fila, c, valor, size=11, fill=VERDE, valign='center',
                   wrap=(c == 3), borde=True)
        fila += 1
    pie = fila + 1
    _celda(ws, pie, 1, PIE, size=8, color=GRIS_SUB)
    ws.merge_cells('A%d:D%d' % (pie, pie))
    _print_setup(ws, header_row=4, landscape=True)
    _area_impresion(ws)
    _proteger(ws, (5, primera_libre + libres - 1))
    return ws


# ==========================================================================
# HELPER 6 · guardar
# ==========================================================================
def guardar(wb, nombre_fichero, titulo, destino=None):
    """Propiedades del libro del contrato §1 + escritura.

    `title` tiene que terminar en « · <sufijo>» o `motor.set_metadata` lo recompone
    (`motor.py:4474-4479`), y `cabeza` tiene que llevar « — <kit>» o `motor.contexto` no sabría
    cómo se llama el kit y el PIE de los 11 ficheros saldría mal (`motor.py:1604-1612`).
    """
    if ' — ' + KIT not in titulo:
        raise SystemExit('El título «%s» no lleva « — %s»: `motor.contexto` no '
                         'podría derivar el nombre del kit ni el pie.' % (titulo, KIT))
    p = wb.properties
    p.creator = CREATOR
    p.lastModifiedBy = CREATOR
    p.title = '%s · %s' % (titulo, SUFIJO)
    p.subject = SUBJECT
    p.keywords = KEYWORDS
    carpeta = destino or OUTPUT_DIR
    if not os.path.isdir(carpeta):
        os.makedirs(carpeta)
    ruta = os.path.join(carpeta, nombre_fichero)
    wb.save(ruta)
    return ruta


# ==========================================================================
# «Se conecta con» — lo que compone motor._bloque_conecta para ESTE kit
# ==========================================================================
F01 = '01-apertura-cierre-taqueria.xlsx'
#: Con la topología de la taquería (no hay fichero de NEGOCIO ni de CAJA: ninguna hoja trae
#: «Denominación | Cantidad» ni «Fecha | Fondo Apertura | Total Facturado»), `papel_del_fichero`
#: da al 01 el papel de ÁREAS (2 ciclos sin «Notas» → 'areas2', `motor.py:4300`) y `_bloque_conecta`
#: entra por la rama ESPEJO (CB-E8) sólo en el 01. El resto cae en la cola genérica.
#: ⚠️ Si al llegar el CONTENIDO real la columna E del 07 pasa a «Cadencia» y la del 08 a
#: «Antelación», el motor asignará `f_periodico` / `f_eventos` y esas dos colas CAMBIAN. Se
#: re-mide con un dry-run y se pega aquí: es la única parte del fichero que no se puede derivar.
DETALLE_AREA = F01 + ' — el mismo día con el DETALLE por área (taquería).'


def conecta(fname):
    """Bloque «Se conecta con» por defecto de un fichero que no es el 01."""
    return ('Se conecta con', [
        DETALLE_AREA,
        'Estás en %s: es una capa MÁS, no una repetición — el marco del día '
        'está en %s.' % (fname, F01),
    ])


CONECTA_01 = ('Dónde encaja este fichero', [
    'Estás en %s: este ES el DETALLE por área del día (taquería) — cómo se '
    'abre y se cierra cada zona. Este kit no trae un fichero aparte de '
    'negocio ni de caja: el marco del día es este mismo.' % F01,
])
CONECTA_BONUS01 = ('Se conecta con', [
    DETALLE_AREA,
    'Estás en BONUS-01-briefing-servicio.xlsx: el briefing abre el turno que '
    'luego se trabaja con %s.' % F01,
])


# ==========================================================================
# Bloques propios de «Instrucciones» (contrato §2)
# ==========================================================================
def usar(vinetas):
    return [('h', 'Cómo usar estas plantillas')] + [('b', v) for v in vinetas]


PERSONALIZAR = [
    ('h', 'Cómo personalizar'),
    ('b', 'Las celdas verdes son editables — cambia responsables y horarios'),
    ('b', 'Ajusta las horas límite a tus horarios reales'),
]

USAR_CHECKLIST = usar([
    'Imprime la hoja correspondiente al turno (apertura o cierre)',
    'El responsable del turno reparte las tareas al equipo',
    'Cada persona marca ✓ cuando completa su tarea y firma',
    'Al final del turno, el encargado verifica y firma al pie',
    'Archiva las hojas firmadas — demuestran control operativo',
])


# ==========================================================================
# ESTRUCTURAS DE CONTENIDO — los redactores de F2 rellenan AQUÍ
# ==========================================================================
# Formato de una hoja de checklist:
#   {'tipo': 'checklist', 'nombre': '<pestaña ≤31 car., sin «/»>',
#    'titulo': '<fila 1>', 'fila2': FILA2_*, 'col_tiempo': 'Hora Límite',
#    'secciones': [('NOMBRE DE SECCIÓN', [(tarea, zona, responsable, 'HH:MM'), …]), …]}
# Reparto de tareas por hoja: SPEC §3.1. Reglas de redacción: SPEC §4.
# ==========================================================================

# --- 01 -------------------------------------------------------------------
SEC_APERTURA = [                                    # EJEMPLO — sustituir en F2
    ('TEMPERATURAS Y EQUIPOS', [
        ('Encender el trompo (asador vertical) y comprobar llama uniforme en todos los quemadores',
         'Cocina', 'Taquero', '10:30'),
        ('Calibrar plancha y comal por zonas (alta y media) y anotar la temperatura de trabajo',
         'Cocina', 'Plancha', '10:45'),
        ('Comprobar que el baño maría de guisados llega a 63 °C antes de abrir',
         'Cocina', 'Guisados', '11:00'),
    ]),
    ('SALSAS Y GUARNICIONES', [                     # EJEMPLO — sustituir en F2
        ('Sacar las salsas crudas a la barra fría y comprobar que se mantienen por debajo de 4 °C',
         'Barra', 'Salsas', '11:15'),
        ('Verificar que el cartel de alérgenos está visible y marca la salsa macha (cacahuete y/o sésamo)',
         'Barra', 'Mostrador', '11:20'),
        ('Reponer cebolla, cilantro y limón en recipientes limpios, nunca sobre el resto anterior',
         'Barra', 'Salsas', '11:25'),
    ]),
]
SEC_CIERRE = [                                      # EJEMPLO — sustituir en F2
    ('TROMPO Y PLANCHA', [
        ('Apagar el trompo, retirar la carne sobrante y registrar su destino: abatida o descartada',
         'Cocina', 'Taquero', '23:30'),
        ('Rascar y engrasar la plancha en caliente; limpiar el comal antes de que enfríe',
         'Cocina', 'Plancha', '23:40'),
        ('Vaciar, lavar y desinfectar la bandeja de grasas del asador vertical',
         'Cocina', 'Taquero', '23:45'),
    ]),
    ('SALSAS Y GUISADOS', [                         # EJEMPLO — sustituir en F2
        ('Descartar las salsas crudas del día y anotar la merma en litros',
         'Barra', 'Salsas', '23:50'),
        ('Abatir el guisado apto, etiquetar con fecha y hora, y descartar el no apto',
         'Cocina', 'Guisados', '23:55'),
        ('Revisar que la hoja de temperaturas del turno queda firmada y archivada',
         'Cocina', 'Encargado', '00:05'),
    ]),
]

# --- BONUS-01 -------------------------------------------------------------
BRIEFING_BLOQUES = [                                # EJEMPLO — sustituir en F2
    ('TROMPO DEL DÍA (SI APLICA)', [
        'Peso montado (kg):', 'Hora de montaje:',
        'Hora límite definida en tu APPCC:', 'Previsión de cortes:',
    ]),
    ('GUISADOS DEL DÍA', [
        'Guisado 1 (hora de entrada / hora límite):',
        'Guisado 2 (hora de entrada / hora límite):',
        'Guisado 3 (hora de entrada / hora límite):',
    ]),
    ('AVISO DE ALÉRGENOS', [
        'Salsa macha en barra (cacahuete y/o sésamo) — ☐ sí ☐ no:',
        'Mole del día (frutos de cáscara) — ☐ sí ☐ no:',
        'Tortilla de harina disponible (gluten) — ☐ sí ☐ no:',
    ]),
    ('SALA Y DELIVERY', [
        'Reservas:', 'Grupos:', 'Previsión de delivery:', 'Personal del turno:',
    ]),
]

# --- BONUS-02 -------------------------------------------------------------
CALENDARIO_FILAS = [                                # EJEMPLO — sustituir en F2
    ('Febrero', '2 Feb — Candelaria',
     'Producción de tamales; comunicarlo en Google Business Profile y en redes', '2 semanas'),
    ('Mayo', '5 Mayo — Cinco de Mayo',
     'Campaña sobre todo para EE. UU.; en México se celebra en Puebla', '1 semana'),
    ('Septiembre', '15-16 Sep — Grito e Independencia',
     'Pico del año fuera de México; reforzar pozole y mezcal', '2 semanas'),
    ('Noviembre', '1-2 Nov — Día de Muertos',
     'Producción de pan de muerto y calabaza en tacha', '2 semanas'),
]


def _sec(nombre, tareas):
    """Azúcar para las secciones de ejemplo de los ficheros 02-08."""
    return (nombre, tareas)


#: EJEMPLO — sustituir en F2. Dos secciones de 3 tareas por hoja, sólo para que el esqueleto
#: sea ejecutable y medible. El reparto real (§3.1 de la SPEC) son 22 hojas y ~300 tareas.
def _demo(a, b, zona, resp, h1, h2):
    return [
        _sec(a, [('%s — tarea de ejemplo %d' % (a.capitalize(), i), zona, resp, h1)
                 for i in (1, 2, 3)]),
        _sec(b, [('%s — tarea de ejemplo %d' % (b.capitalize(), i), zona, resp, h2)
                 for i in (1, 2, 3)]),
    ]


FICHEROS = [
    {
        'fichero': F01,
        'titulo': 'Tareas de Apertura y Cierre — ' + KIT,
        'bloques': USAR_CHECKLIST + PERSONALIZAR,
        'conecta': CONECTA_01,
        'hojas': [
            {'tipo': 'checklist', 'nombre': 'Apertura Taquería',
             'titulo': 'Checklist de Apertura — Taquería',
             'fila2': FILA2_TURNO, 'col_tiempo': 'Hora Límite',
             'secciones': SEC_APERTURA},
            {'tipo': 'checklist', 'nombre': 'Cierre Taquería',
             'titulo': 'Checklist de Cierre — Taquería',
             'fila2': FILA2_TURNO, 'col_tiempo': 'Hora Límite',
             'secciones': SEC_CIERRE},
        ],
    },
    {
        'fichero': '02-trompo-plancha-comal.xlsx',
        'titulo': 'Trompo, Plancha y Comal — ' + KIT,
        'bloques': USAR_CHECKLIST + PERSONALIZAR,
        'conecta': conecta('02-trompo-plancha-comal.xlsx'),
        'hojas': [
            {'tipo': 'checklist', 'nombre': 'Marinado y Montaje del Trompo',
             'fila2': FILA2_TURNO, 'col_tiempo': 'Hora Límite',
             'secciones': _demo('MARINADO (VÍSPERA)', 'MONTAJE DEL TROMPO (SI APLICA)',
                                'Cocina', 'Taquero', '09:00', '10:30')},
            {'tipo': 'checklist', 'nombre': 'Corte, Plancha y Comal',
             'fila2': FILA2_TURNO, 'col_tiempo': 'Hora Límite',
             'secciones': _demo('CORTE Y SONDA', 'PLANCHA Y COMAL',
                                'Cocina', 'Taquero', '13:00', '13:30')},
        ],
    },
    {
        # ⚠️ El TÍTULO del documento no puede decir «APPCC»: `motor.texto_appcc` le añadiría la
        # coletilla del Pack APPCC a la celda B2 de «Instrucciones», de donde sale el título del
        # libro. El NOMBRE DE FICHERO sí puede («appcc» en minúscula no casa con `RX_APPCC`).
        'fichero': '03-appcc-salsas-alergenos.xlsx',
        'titulo': 'Temperaturas, Salsas y Alérgenos — ' + KIT,
        'bloques': USAR_CHECKLIST + PERSONALIZAR,
        'conecta': conecta('03-appcc-salsas-alergenos.xlsx'),
        'hojas': [
            {'tipo': 'checklist', 'nombre': 'Temperaturas y Trazabilidad',
             'fila2': FILA2_TURNO, 'col_tiempo': 'Hora Límite',
             'secciones': _demo('TEMPERATURAS (2 VECES POR TURNO)', 'RECEPCIÓN Y TRAZABILIDAD',
                                'Cocina', 'Encargado', '11:00', '12:00')},
            {'tipo': 'checklist', 'nombre': 'Barra de Salsas',
             'fila2': FILA2_TURNO, 'col_tiempo': 'Hora Límite',
             'secciones': _demo('MONTAJE DE LA BARRA', 'CONTROL DURANTE EL SERVICIO',
                                'Barra', 'Salsas', '11:15', '15:00')},
            {'tipo': 'checklist', 'nombre': 'Alérgenos y Cruzada',
             'fila2': FILA2_TURNO, 'col_tiempo': 'Hora Límite',
             'secciones': _demo('CACAHUETE Y SÉSAMO', 'GLUTEN: MAÍZ FRENTE A HARINA',
                                'Barra', 'Mostrador', '11:20', '11:30')},
        ],
    },
    {
        'fichero': '04-nixtamal-tortilla-guisados.xlsx',
        'titulo': 'Nixtamal, Tortilla y Guisados — ' + KIT,
        'bloques': USAR_CHECKLIST + PERSONALIZAR,
        'conecta': conecta('04-nixtamal-tortilla-guisados.xlsx'),
        'hojas': [
            {'tipo': 'checklist', 'nombre': 'Nixtamal y Molienda',
             'fila2': FILA2_TURNO, 'col_tiempo': 'Hora Límite',
             'secciones': _demo('COCCIÓN Y REPOSO (SI APLICA)', 'MOLIENDA Y MASA',
                                'Tortillería', 'Tortillería', '06:00', '07:30')},
            {'tipo': 'checklist', 'nombre': 'Tortilla y Comal',
             'fila2': FILA2_TURNO, 'col_tiempo': 'Hora Límite',
             'secciones': _demo('PRENSADO Y COMAL', 'MERMAS Y LIMPIEZA',
                                'Tortillería', 'Tortillería', '08:00', '23:30')},
            {'tipo': 'checklist', 'nombre': 'Guisados y Rotación',
             'fila2': FILA2_TURNO, 'col_tiempo': 'Hora Límite',
             'secciones': _demo('ELABORACIÓN', 'ROTACIÓN Y BAÑO MARÍA',
                                'Cocina', 'Guisados', '09:30', '12:00')},
        ],
    },
    {
        'fichero': '05-tareas-manager.xlsx',
        'titulo': 'Tareas del Manager — ' + KIT,
        'bloques': USAR_CHECKLIST + PERSONALIZAR,
        'conecta': conecta('05-tareas-manager.xlsx'),
        'hojas': [
            {'tipo': 'checklist', 'nombre': 'Tareas Diarias Manager',
             'fila2': FILA2_TURNO, 'col_tiempo': 'Hora Límite',
             'secciones': _demo('ANTES DE APERTURA', 'CIERRE DE NEGOCIO',
                                'Oficina', 'Manager', '10:00', '00:15')},
            # ⚠️ En F2, con la columna E llena de días, hay que pasar
            # col_tiempo='Día' y fila2=FILA2_DIA o el motor los reescribe.
            {'tipo': 'checklist', 'nombre': 'Tareas Semanales Manager',
             'fila2': FILA2_TURNO, 'col_tiempo': 'Hora Límite',
             'secciones': _demo('EQUIPO Y TURNOS', 'PROVEEDORES Y COMPRAS',
                                'Oficina', 'Manager', '10:30', '11:00')},
        ],
    },
    {
        'fichero': '06-tareas-perfiles.xlsx',
        'titulo': 'Tareas por Perfil — ' + KIT,
        'bloques': USAR_CHECKLIST + PERSONALIZAR,
        'conecta': conecta('06-tareas-perfiles.xlsx'),
        'hojas': [
            {'tipo': 'checklist', 'nombre': nombre, 'fila2': FILA2_TURNO,
             'col_tiempo': 'Hora Límite',
             'secciones': _demo('ANTES DEL SERVICIO', 'DURANTE Y AL CIERRE',
                                zona, nombre, h1, h2)}
            for nombre, zona, h1, h2 in (
                ('Taquero del Trompo', 'Cocina', '10:30', '23:30'),
                ('Tortillería', 'Tortillería', '06:00', '23:30'),
                ('Salsas', 'Barra', '09:00', '23:50'),
                ('Plancha y Freidora', 'Cocina', '10:45', '23:40'),
                ('Mostrador y Caja', 'Sala', '11:30', '00:05'),
                ('Reparto y Delivery', 'Sala', '12:00', '23:45'),
            )
        ],
    },
    {
        'fichero': '07-semanales-mensuales.xlsx',
        'titulo': 'Tareas Semanales y Mensuales — ' + KIT,
        'bloques': USAR_CHECKLIST + PERSONALIZAR,
        'conecta': conecta('07-semanales-mensuales.xlsx'),
        'hojas': [
            {'tipo': 'checklist', 'nombre': 'Tareas Semanales',
             'fila2': FILA2_TURNO, 'col_tiempo': 'Hora Límite',
             'secciones': _demo('LIMPIEZA PROFUNDA', 'INVENTARIO Y ALMACÉN',
                                'Cocina', 'Encargado', '09:00', '10:00')},
            {'tipo': 'checklist', 'nombre': 'Tareas Mensuales',
             'fila2': FILA2_TURNO, 'col_tiempo': 'Hora Límite',
             'secciones': _demo('MANTENIMIENTO', 'DOCUMENTACIÓN',
                                'Cocina', 'Manager', '09:30', '11:00')},
        ],
    },
    {
        'fichero': '08-eventos-estacionales.xlsx',
        'titulo': 'Eventos y Temporadas — ' + KIT,
        'bloques': USAR_CHECKLIST + PERSONALIZAR,
        'conecta': conecta('08-eventos-estacionales.xlsx'),
        'hojas': [
            {'tipo': 'checklist', 'nombre': 'Temporadas y Producto',
             'fila2': FILA2_ANUAL, 'col_tiempo': 'Hora Límite',
             'secciones': _demo('CHILE SECO POR CAMPAÑA', 'CARNE Y MAÍZ',
                                'Almacén', 'Manager', '09:00', '09:30')},
            {'tipo': 'checklist', 'nombre': 'Calendario de Eventos',
             'fila2': FILA2_ANUAL, 'col_tiempo': 'Hora Límite',
             'secciones': _demo('PRIMER TRIMESTRE', 'CUARTO TRIMESTRE',
                                'Todos', 'Manager', '09:00', '09:30')},
        ],
    },
    {
        'fichero': '09-plantilla-personalizable.xlsx',
        'titulo': 'Plantilla Personalizable — ' + KIT,
        'bloques': usar([
            'Escribe tus tareas en las celdas verdes de la columna «Tarea»',
            'Renombra las secciones con los bloques reales de tu taquería',
            'El contador arranca en 0 y sube a medida que escribes: cuenta lo que hay',
        ]) + PERSONALIZAR,
        'conecta': conecta('09-plantilla-personalizable.xlsx'),
        'hojas': [
            {'tipo': 'plantilla', 'nombre': 'Plantilla en Blanco',
             'titulo': 'Plantilla en Blanco — ' + KIT,
             'fila2': FILA2_TURNO, 'col_tiempo': 'Hora Límite',
             'secciones': ['SECCIÓN 1 (PERSONALIZAR)', 'SECCIÓN 2 (PERSONALIZAR)',
                           'SECCIÓN 3 (PERSONALIZAR)']},
        ],
    },
    {
        'fichero': 'BONUS-01-briefing-servicio.xlsx',
        'titulo': 'Briefing Pre-Servicio — ' + KIT,
        'bloques': usar([
            'Imprime una hoja por turno y rellénala antes del pase',
            'El encargado la repasa con el equipo en dos minutos y firma al pie',
            'Archívala con la hoja de temperaturas del mismo turno',
        ]) + [
            ('h', 'Cómo personalizar'),
            ('b', 'Las celdas verdes son editables — escribe ahí los datos del turno'),
            ('b', 'Añade o quita líneas de guisado y de salsa según tu carta'),
        ],
        'conecta': CONECTA_BONUS01,
        'con_checklist': False,
        'hojas': [{'tipo': 'briefing', 'bloques': BRIEFING_BLOQUES}],
    },
    {
        'fichero': 'BONUS-02-calendario-anual.xlsx',
        'titulo': 'Calendario Anual de Tareas — ' + KIT,
        'bloques': usar([
            'Imprímelo en A4 apaisado y cuélgalo en la oficina',
            'Cada fila es una fecha o una tarea anual con su antelación',
            'Las filas verdes del final son para tus fechas locales',
        ]) + [
            ('h', 'Cómo personalizar'),
            ('b', 'Las celdas verdes son editables — escribe ahí tus fechas'),
            ('b', 'Ajusta la antelación a los plazos de tus proveedores'),
        ],
        'conecta': conecta('BONUS-02-calendario-anual.xlsx'),
        'con_checklist': False,
        'hojas': [{'tipo': 'calendario', 'filas': CALENDARIO_FILAS,
                   'titulo': 'Calendario Anual — Fechas Clave de la Taquería',
                   'fila2': FILA2_ANUAL}],
    },
]


# ==========================================================================
# Construcción
# ==========================================================================
# ==========================================================================
# CONTENIDO REAL (F2): los redactores escriben `contenido/contenido_a.py` (01-04) y
# `contenido_b.py` (05-09 y BONUS) sin tocar este fichero. Si existen, sustituyen los
# EJEMPLOS de arriba por nombre de hoja; si no, el esqueleto sigue siendo ejecutable.
# La columna de tiempo de 07/08 se fija aquí para que coincida con lo que el motor
# decidiría por CONTENIDO (motor.cadencia): día de la semana → «Día», cadencia →
# «Cadencia», antelación → «Antelación». Verificado con main.py --dry-run.
# ==========================================================================
COL_TIEMPO_POR_HOJA = {
    'Tareas Semanales': ('Día', FILA2_DIA),
    'Tareas Mensuales': ('Cadencia', FILA2_CADENCIA),
    'Temporadas y Producto': ('Antelación', FILA2_ANTELACION),
    'Calendario de Eventos': ('Antelación', FILA2_ANTELACION),
}


def aplicar_contenido(ficheros):
    import importlib
    base = os.path.join(os.path.dirname(os.path.abspath(__file__)),
                        'productos-digitales', 'kit-tareas-taqueria', 'contenido')
    if base not in sys.path:
        sys.path.insert(0, base)
    mods = []
    for nombre in ('contenido_a', 'contenido_b'):
        try:
            mods.append(importlib.import_module(nombre))
        except ImportError:
            pass
    sec = {}
    for m in mods:
        sec.update(getattr(m, 'SEC', {}))
    usados = set()
    for cfg in ficheros:
        for h in cfg['hojas']:
            tipo = h['tipo']
            if tipo == 'checklist' and h['nombre'] in sec:
                h['secciones'] = sec[h['nombre']]
                usados.add(h['nombre'])
            if tipo == 'checklist' and h['nombre'] in COL_TIEMPO_POR_HOJA:
                h['col_tiempo'], h['fila2'] = COL_TIEMPO_POR_HOJA[h['nombre']]
            for m in mods:
                if tipo == 'plantilla' and hasattr(m, 'PLANTILLA_SECCIONES'):
                    h['secciones'] = m.PLANTILLA_SECCIONES
                if tipo == 'briefing' and hasattr(m, 'BRIEFING_BLOQUES'):
                    h['bloques'] = m.BRIEFING_BLOQUES
                if tipo == 'calendario' and hasattr(m, 'CALENDARIO_FILAS'):
                    h['filas'] = m.CALENDARIO_FILAS
    sobran = set(sec) - usados
    if sobran:
        raise SystemExit('Hojas del contenido que no existen en FICHEROS: %s' % sorted(sobran))
    return len(mods), len(usados)


N_MODULOS, N_HOJAS_REALES = aplicar_contenido(FICHEROS)


def construir(cfg, destino, prueba=False):
    wb = Workbook()
    wb.active.title = 'Instrucciones'
    hoja_instrucciones(wb, cfg['titulo'], cfg['bloques'], cfg['conecta'],
                       con_checklist=cfg.get('con_checklist', True))
    for h in cfg['hojas']:
        tipo = h['tipo']
        if tipo == 'checklist':
            secciones = h['secciones']
            if prueba:
                secciones = [(n, t[:3]) for n, t in secciones]
            hoja_checklist(wb, h['nombre'], h['fila2'], secciones,
                           col_tiempo=h.get('col_tiempo', 'Hora Límite'),
                           titulo=h.get('titulo'))
        elif tipo == 'plantilla':
            hoja_plantilla(wb, h['nombre'], h['fila2'], h['secciones'],
                           col_tiempo=h.get('col_tiempo', 'Hora Límite'),
                           titulo=h.get('titulo'))
        elif tipo == 'briefing':
            hoja_briefing(wb, h['bloques'])
        elif tipo == 'calendario':
            hoja_calendario(wb, h['filas'], titulo=h.get('titulo'),
                            fila2=h.get('fila2'))
        else:
            raise SystemExit('Tipo de hoja desconocido: %s' % tipo)
    return guardar(wb, cfg['fichero'], cfg['titulo'], destino)


def main():
    ap = argparse.ArgumentParser(description='Generador de %s' % PID)
    ap.add_argument('--dir', default=None,
                    help='carpeta de salida (por defecto astro-site/public/dl/%s)' % PID)
    ap.add_argument('--prueba', action='store_true',
                    help='sólo el 01, con 3 tareas por sección; EXIGE --dir '
                         '(no toca astro-site/public/dl/)')
    args = ap.parse_args()
    if args.prueba and not args.dir:
        raise SystemExit('--prueba exige --dir (un directorio del scratchpad): el modo de '
                         'prueba NO escribe en astro-site/public/dl/.')
    destino = os.path.abspath(args.dir) if args.dir else OUTPUT_DIR
    ficheros = FICHEROS[:1] if args.prueba else FICHEROS
    for cfg in ficheros:
        ruta = construir(cfg, destino, prueba=args.prueba)
        print('  %s  (%d hojas)' % (os.path.basename(ruta), len(cfg['hojas']) + 1))
    print('%d ficheros en %s' % (len(ficheros), destino))
    return 0


if __name__ == '__main__':
    sys.exit(main())
